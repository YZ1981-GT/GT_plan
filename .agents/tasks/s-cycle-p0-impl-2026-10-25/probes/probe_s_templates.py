"""Probe: list sheet names (and first header-ish rows) of S1/S8/S9/S10/S17 templates.

S17 is .xls -> xlrd; others .xlsx -> openpyxl. Read-only; prints to stdout as UTF-8.
"""
from __future__ import annotations

import sys
from pathlib import Path

sys.stdout.reconfigure(encoding="utf-8")

ROOT = Path(__file__).resolve().parents[4]
TPL = ROOT / "backend" / "wp_templates" / "S"

TARGETS = {
    "S1": "S1 对被审计单位违反法规行为的考虑2019.xlsx",
    "S8": "S8 租赁.xlsx",
    "S9": "S9 对电子商务的考虑.xlsx",
    "S10": "S10 对环境事项的考虑.xlsx",
    "S17": "S17 非经常性损益.xls",
}


def _row_text(values: list) -> str:
    cells = [str(v).strip().replace("\n", " ") for v in values if v not in (None, "")]
    return " | ".join(c[:24] for c in cells)[:220]


def dump_xlsx(path: Path) -> None:
    import openpyxl

    wb = openpyxl.load_workbook(path, read_only=True, data_only=True)
    for ws in wb.worksheets:
        print(f"  [sheet] {ws.title!r}  dims={ws.max_row}x{ws.max_column}")
        shown = 0
        for row in ws.iter_rows(min_row=1, max_row=12, values_only=True):
            text = _row_text(list(row))
            if text:
                print(f"      {text}")
                shown += 1
            if shown >= 6:
                break
    wb.close()


def dump_xls(path: Path) -> None:
    import xlrd

    book = xlrd.open_workbook(str(path))
    for sh in book.sheets():
        print(f"  [sheet] {sh.name!r}  dims={sh.nrows}x{sh.ncols}")
        shown = 0
        for r in range(min(sh.nrows, 12)):
            text = _row_text(sh.row_values(r))
            if text:
                print(f"      {text}")
                shown += 1
            if shown >= 6:
                break


def main() -> None:
    for code, name in TARGETS.items():
        path = TPL / name
        print(f"== {code}: {name} exists={path.exists()}")
        if not path.exists():
            continue
        try:
            if path.suffix.lower() == ".xls":
                dump_xls(path)
            else:
                dump_xlsx(path)
        except Exception as exc:  # noqa: BLE001 - probe reports, never raises
            print(f"  !! read failed: {type(exc).__name__}: {exc}")


if __name__ == "__main__":
    main()

"""M5 盈余公积 — 业务逻辑服务

提供功能：
- 导出模板 / 导出数据 / 导入数据（xlsx三级）
- 法定计提测试（纯函数）
- 50%上限判断

科目4101盈余公积（**贷方/权益类！**）：期末=期初+贷方-借方
法定盈余公积按净利润10%计提，累计达注册资本50%可停止计提。
任意盈余公积由股东大会决议自主计提。

Requirements: 4.3, 6.1, 6.6
"""

from __future__ import annotations

import io
import json
import logging
from typing import Any

from openpyxl import Workbook, load_workbook
from openpyxl.styles import Alignment, Font, PatternFill
from sqlalchemy.ext.asyncio import AsyncSession

logger = logging.getLogger(__name__)

_M5_ACCOUNT_CODE = "4101"
_ROW_LIMIT = 500


# ═══════════════════════════════════════════════════════════════════════════════
# 纯函数：计提计算引擎
# ═══════════════════════════════════════════════════════════════════════════════


def calc_statutory_accrual(base: float, rate: float = 0.1) -> float:
    """法定盈余公积应计提 = 计提基数 × 比例（默认10%）"""
    return base * rate


def calc_accrual_diff(estimated: float, booked: float) -> float:
    """计提差异 = 应计提 − 账面计提"""
    return estimated - booked


def is_ceiling_reached(accumulated: float, registered_capital: float) -> bool:
    """累计法定盈余公积是否达注册资本50%"""
    if registered_capital <= 0:
        return False
    return accumulated >= registered_capital * 0.5


def run_accrual_test(
    *,
    net_profit: float,
    prior_loss_offset: float,
    statutory_rate: float = 0.1,
    discretionary_rate: float = 0.0,
    statutory_booked: float = 0.0,
    discretionary_booked: float = 0.0,
    accumulated_statutory: float = 0.0,
    registered_capital: float = 0.0,
) -> dict[str, Any]:
    """执行完整计提测试，返回结果字典

    计算逻辑：
    1. 计提基数 = 净利润 - 弥补以前年度亏损
    2. 法定应计提 = 基数 × statutory_rate
    3. 任意应计提 = 基数 × discretionary_rate
    4. 差异 = 应计提 - 账面
    5. 50%上限判断
    """
    accrual_base = net_profit - prior_loss_offset

    statutory_estimated = calc_statutory_accrual(accrual_base, statutory_rate)
    statutory_diff = calc_accrual_diff(statutory_estimated, statutory_booked)

    discretionary_estimated = calc_statutory_accrual(accrual_base, discretionary_rate)
    discretionary_diff = calc_accrual_diff(discretionary_estimated, discretionary_booked)

    total_estimated = statutory_estimated + discretionary_estimated
    total_diff = calc_accrual_diff(total_estimated, statutory_booked + discretionary_booked)

    ceiling_reached = is_ceiling_reached(accumulated_statutory, registered_capital)

    return {
        "accrual_base": accrual_base,
        "statutory_estimated": statutory_estimated,
        "statutory_diff": statutory_diff,
        "discretionary_estimated": discretionary_estimated,
        "discretionary_diff": discretionary_diff,
        "total_estimated": total_estimated,
        "total_diff": total_diff,
        "ceiling_reached": ceiling_reached,
    }


# ═══════════════════════════════════════════════════════════════════════════════
# 导入导出配置
# ═══════════════════════════════════════════════════════════════════════════════

_SHEET_CONFIGS: dict[str, dict[str, Any]] = {
    "M5-2": {
        "title": "M5-2 明细表（法定+任意盈余公积）",
        "headers": [
            "区段", "来源项目", "期初余额", "本期计提",
            "本期转增资本", "本期弥补亏损", "期末余额", "备注",
        ],
        "fields": [
            "segment", "sourceName", "beginning", "accrual",
            "capitalConversion", "lossOffset", "endBalance", "remark",
        ],
    },
}

_HEADER_FONT = Font(bold=True, size=11)
_HEADER_FILL = PatternFill(start_color="4472C4", end_color="4472C4", fill_type="solid")
_HEADER_FONT_WHITE = Font(bold=True, size=11, color="FFFFFF")
_HEADER_ALIGN = Alignment(horizontal="center", vertical="center", wrap_text=True)


# ═══════════════════════════════════════════════════════════════════════════════
# 导出模板
# ═══════════════════════════════════════════════════════════════════════════════


# ═══════════════════════════════════════════════════════════════════════════════
# 跨底稿联动：从 M6 获取净利润/计提基数
# ═══════════════════════════════════════════════════════════════════════════════


async def fetch_m6_net_profit(
    project_id: str,
    db: AsyncSession,
) -> dict[str, Any]:
    """从M6未分配利润底稿获取净利润/计提基数

    查询逻辑：
    1. 通过 wp_index 找到同项目下 wp_code='M6' 的底稿
    2. 从该底稿的 checklist_responses 中读取净利润相关 item_id
    3. 返回 { net_profit, prior_loss_offset, accrual_base, source_wp_id, ready }

    联动方向：M6未分配利润 → M5计提检查表（Req 4.2）
    """
    import sqlalchemy as sa
    from app.models.audit_platform_models import (
        ChecklistResponse,
        WorkingPaper,
        WpIndex,
    )

    # Step 1: 找 M6 底稿
    stmt = (
        sa.select(WorkingPaper.id)
        .join(WpIndex, WpIndex.wp_id == WorkingPaper.id)
        .where(
            WorkingPaper.project_id == project_id,
            WpIndex.wp_code == "M6",
        )
        .limit(1)
    )
    result = await db.execute(stmt)
    m6_wp_id = result.scalar_one_or_none()

    if not m6_wp_id:
        return {
            "net_profit": 0.0,
            "prior_loss_offset": 0.0,
            "accrual_base": 0.0,
            "source_wp_id": None,
            "ready": False,
            "message": "M6未分配利润底稿未找到",
        }

    # Step 2: 从 M6 的 checklist_responses 读取净利润数据
    # M6 存储约定: item_id "M6-net-profit" (remark=净利润金额)
    #              item_id "M6-prior-loss-offset" (remark=弥补以前年度亏损)
    stmt = sa.select(ChecklistResponse).where(
        ChecklistResponse.wp_id == m6_wp_id,
        ChecklistResponse.item_id.in_([
            "M6-net-profit",
            "M6-1-net-profit",
            "M6-prior-loss-offset",
            "M6-1-prior-loss-offset",
        ]),
    )
    result = await db.execute(stmt)
    responses = result.scalars().all()

    net_profit = 0.0
    prior_loss_offset = 0.0

    for resp in responses:
        try:
            val = float(resp.remark) if resp.remark else 0.0
        except (ValueError, TypeError):
            val = 0.0

        if "net-profit" in resp.item_id:
            net_profit = val
        elif "prior-loss-offset" in resp.item_id:
            prior_loss_offset = val

    accrual_base = net_profit - prior_loss_offset

    return {
        "net_profit": net_profit,
        "prior_loss_offset": prior_loss_offset,
        "accrual_base": accrual_base,
        "source_wp_id": m6_wp_id,
        "ready": net_profit != 0.0,
        "message": "M6净利润已就绪" if net_profit != 0.0 else "M6净利润数据为空，请先完成M6底稿",
    }


async def generate_export_template(
    wp_id: str,
    db: AsyncSession,
    *,
    sheet: str | None = None,
) -> io.BytesIO:
    """导出空白xlsx模板

    Args:
        wp_id: 底稿ID
        db: 数据库会话
        sheet: 可选，指定导出单个sheet（M5-2）
    """
    wb = Workbook()
    wb.remove(wb.active)  # type: ignore[arg-type]

    sheets_to_export = [sheet] if sheet and sheet in _SHEET_CONFIGS else list(_SHEET_CONFIGS.keys())

    for sheet_key in sheets_to_export:
        cfg = _SHEET_CONFIGS[sheet_key]
        ws = wb.create_sheet(title=cfg["title"])

        # 写入表头
        for col_idx, header in enumerate(cfg["headers"], 1):
            cell = ws.cell(row=1, column=col_idx, value=header)
            cell.font = _HEADER_FONT_WHITE
            cell.fill = _HEADER_FILL
            cell.alignment = _HEADER_ALIGN

        # 设置列宽
        col_widths = [12, 30, 15, 15, 15, 15, 15, 25]
        for col_idx, width in enumerate(col_widths, 1):
            ws.column_dimensions[chr(64 + col_idx)].width = width

    buffer = io.BytesIO()
    wb.save(buffer)
    buffer.seek(0)
    return buffer


# ═══════════════════════════════════════════════════════════════════════════════
# 导出数据
# ═══════════════════════════════════════════════════════════════════════════════


async def generate_export_data(
    wp_id: str,
    db: AsyncSession,
    *,
    sheet: str | None = None,
) -> io.BytesIO:
    """导出含当前数据的xlsx

    从 checklist_responses 查询 M5-2 行数据，写入 xlsx。
    """
    from app.models.audit_platform_models import ChecklistResponse
    import sqlalchemy as sa

    wb = Workbook()
    wb.remove(wb.active)  # type: ignore[arg-type]

    sheets_to_export = [sheet] if sheet and sheet in _SHEET_CONFIGS else list(_SHEET_CONFIGS.keys())

    for sheet_key in sheets_to_export:
        cfg = _SHEET_CONFIGS[sheet_key]
        ws = wb.create_sheet(title=cfg["title"])

        # 写入表头
        for col_idx, header in enumerate(cfg["headers"], 1):
            cell = ws.cell(row=1, column=col_idx, value=header)
            cell.font = _HEADER_FONT_WHITE
            cell.fill = _HEADER_FILL
            cell.alignment = _HEADER_ALIGN

        # 查询 checklist_responses 中的行数据
        stmt = sa.select(ChecklistResponse).where(
            ChecklistResponse.wp_id == wp_id,
            ChecklistResponse.item_id.like(f"M5-2-row-%-data"),
        ).order_by(ChecklistResponse.item_id)

        result = await db.execute(stmt)
        responses = result.scalars().all()

        row_num = 2
        for resp in responses:
            if not resp.remark:
                continue
            try:
                data = json.loads(resp.remark)
            except (json.JSONDecodeError, TypeError):
                continue

            for col_idx, field in enumerate(cfg["fields"], 1):
                value = data.get(field, "")
                ws.cell(row=row_num, column=col_idx, value=value)
            row_num += 1

    buffer = io.BytesIO()
    wb.save(buffer)
    buffer.seek(0)
    return buffer


# ═══════════════════════════════════════════════════════════════════════════════
# 导入数据
# ═══════════════════════════════════════════════════════════════════════════════


async def import_data(
    wp_id: str,
    content: bytes,
    db: AsyncSession,
    *,
    sheet: str | None = None,
) -> dict[str, Any]:
    """导入xlsx解析写入checklist_responses

    解析xlsx文件，按 M5-2 明细表结构写入。
    权益类贷方：期末=期初+计提-转增-弥补

    Returns:
        { imported_count: int, warnings: list[str] }
    """
    from app.models.audit_platform_models import ChecklistResponse
    import sqlalchemy as sa

    target_sheet = sheet or "M5-2"
    if target_sheet not in _SHEET_CONFIGS:
        raise ValueError(f"不支持的sheet: {target_sheet}")

    cfg = _SHEET_CONFIGS[target_sheet]
    expected_headers = cfg["headers"]
    fields = cfg["fields"]

    try:
        wb = load_workbook(io.BytesIO(content), read_only=True, data_only=True)
    except Exception as e:
        raise ValueError(f"无法解析xlsx文件: {e}")

    # 找到目标工作表（优先按标题匹配，其次取第一个）
    ws = None
    for ws_name in wb.sheetnames:
        if cfg["title"] in ws_name or target_sheet in ws_name:
            ws = wb[ws_name]
            break
    if ws is None:
        ws = wb.active

    if ws is None:
        raise ValueError("xlsx文件为空")

    # 验证表头
    header_row = [cell.value for cell in next(ws.iter_rows(min_row=1, max_row=1))]
    # 宽容匹配：只要包含关键列即可
    missing = [h for h in expected_headers[:6] if h not in header_row]
    if missing:
        raise ValueError(f"列头不匹配，缺少: {missing}")

    # 解析数据行
    warnings: list[str] = []
    imported_count = 0
    items_to_save: list[dict[str, Any]] = []

    for row_idx, row in enumerate(ws.iter_rows(min_row=2, values_only=True), start=1):
        if row_idx > _ROW_LIMIT:
            warnings.append(f"超过最大行限制({_ROW_LIMIT})，后续行已忽略")
            break

        # 跳过全空行
        if all(cell is None or cell == "" for cell in row):
            continue

        row_data: dict[str, Any] = {}
        for col_idx, field in enumerate(fields):
            val = row[col_idx] if col_idx < len(row) else None
            row_data[field] = val if val is not None else ""

        # 计算期末（权益类贷方：期末=期初+计提-转增-弥补）
        try:
            beginning = float(row_data.get("beginning") or 0)
            accrual = float(row_data.get("accrual") or 0)
            conversion = float(row_data.get("capitalConversion") or 0)
            loss = float(row_data.get("lossOffset") or 0)
            row_data["endBalance"] = beginning + accrual - conversion - loss
        except (ValueError, TypeError):
            pass

        items_to_save.append(row_data)
        imported_count += 1

    # 写入 checklist_responses
    for idx, row_data in enumerate(items_to_save, start=1):
        item_id = f"M5-2-row-{idx}-data"

        # Upsert: 查询是否存在
        stmt = sa.select(ChecklistResponse).where(
            ChecklistResponse.wp_id == wp_id,
            ChecklistResponse.item_id == item_id,
        )
        result = await db.execute(stmt)
        existing = result.scalar_one_or_none()

        remark_json = json.dumps(row_data, ensure_ascii=False)

        if existing:
            existing.remark = remark_json
        else:
            new_resp = ChecklistResponse(
                wp_id=wp_id,
                item_id=item_id,
                conclusion=None,
                remark=remark_json,
            )
            db.add(new_resp)

    return {"imported_count": imported_count, "warnings": warnings}

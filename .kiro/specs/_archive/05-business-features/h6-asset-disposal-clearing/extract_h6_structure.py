"""
Phase 0 Task 0.1: openpyxl脚本读取H6固定资产清理.xlsx全部8 sheet
提取：sheet名/列头/行数/公式单元格/合并区域/数据类型
产出：h6_structure_summary.json
验证：8 sheet结构与spec描述一致
"""

import json
import os
from pathlib import Path

from openpyxl import load_workbook
from openpyxl.utils import get_column_letter

# Paths
XLSX_PATH = Path(r"d:\GT_plan\基础数据\致同通用审计程序及底稿模板（2025年修订）\1.致同审计程序及底稿模板（2025年）\4.风险应对-实质性程序（D-N）\H 固定资产循环\H6 固定资产清理.xlsx")
OUTPUT_PATH = Path(r"d:\GT_plan\.kiro\specs\h6-asset-disposal-clearing\h6_structure_summary.json")

# Expected sheets from spec Glossary
EXPECTED_SHEETS = [
    "Tab_Index",
    "Procedure_Table_H6A",
    "Adjudication_H6_1",
    "Disclosure_Listed",
    "Disclosure_SOE",
    "Detail_H6_2",
    "Adjustment_H6_3",
    "Check_H6_4",
]

# Spec expected structure (for validation)
SPEC_EXPECTED = {
    "Tab_Index": {"rows": 19, "cols": 8},
    "Procedure_Table_H6A": {"rows": 23, "cols": 12},
    "Adjudication_H6_1": {"rows": 51, "cols": 9, "formulas": 59},
    "Disclosure_Listed": {"rows": 18, "cols": 5},
    "Disclosure_SOE": {"rows": 18, "cols": 5},
    "Detail_H6_2": {"rows": 36, "cols": 25, "formulas": 16},
    "Adjustment_H6_3": {"rows": 21, "cols": 10},
    "Check_H6_4": {"rows": 34, "cols": 18, "formulas": 8},
}


def extract_sheet_info(ws):
    """Extract comprehensive info from a worksheet."""
    info = {
        "sheet_name": ws.title,
        "dimensions": ws.dimensions,
        "max_row": ws.max_row,
        "max_column": ws.max_column,
        "column_headers": [],
        "formula_cells": [],
        "merged_areas": [],
        "data_types": {},
        "sample_data": [],
    }

    # Extract column headers (row 1)
    headers = []
    for col in range(1, ws.max_column + 1):
        cell = ws.cell(row=1, column=col)
        val = cell.value
        headers.append({
            "column": get_column_letter(col),
            "col_index": col,
            "value": str(val) if val is not None else None,
            "data_type": cell.data_type,
        })
    info["column_headers"] = headers

    # Also check row 2 for headers (some sheets have merged header rows)
    row2_headers = []
    if ws.max_row >= 2:
        for col in range(1, ws.max_column + 1):
            cell = ws.cell(row=2, column=col)
            val = cell.value
            if val is not None:
                row2_headers.append({
                    "column": get_column_letter(col),
                    "col_index": col,
                    "value": str(val),
                    "data_type": cell.data_type,
                })
    if row2_headers:
        info["row2_headers"] = row2_headers

    # Extract formula cells
    formula_cells = []
    for row in ws.iter_rows(min_row=1, max_row=ws.max_row, min_col=1, max_col=ws.max_column):
        for cell in row:
            if cell.data_type == 'f':
                formula_cells.append({
                    "cell": f"{get_column_letter(cell.column)}{cell.row}",
                    "row": cell.row,
                    "column": cell.column,
                    "formula": str(cell.value),
                })
    info["formula_cells"] = formula_cells
    info["formula_count"] = len(formula_cells)

    # Extract merged areas
    merged = []
    for merged_range in ws.merged_cells.ranges:
        merged.append(str(merged_range))
    info["merged_areas"] = merged
    info["merged_count"] = len(merged)

    # Data type distribution
    type_counts = {}
    for row in ws.iter_rows(min_row=1, max_row=ws.max_row, min_col=1, max_col=ws.max_column):
        for cell in row:
            dt = cell.data_type
            type_counts[dt] = type_counts.get(dt, 0) + 1
    # Map data_type codes to human-readable
    type_map = {
        's': 'string',
        'n': 'numeric',
        'f': 'formula',
        'b': 'boolean',
        'd': 'date',
        'e': 'error',
        None: 'empty',
    }
    info["data_types"] = {type_map.get(k, k): v for k, v in type_counts.items()}

    # Sample data: first 5 rows
    sample = []
    for row_idx, row in enumerate(ws.iter_rows(min_row=1, max_row=min(5, ws.max_row), min_col=1, max_col=ws.max_column), start=1):
        row_data = {}
        for cell in row:
            if cell.value is not None:
                col_letter = get_column_letter(cell.column)
                row_data[col_letter] = {
                    "value": str(cell.value),
                    "type": cell.data_type,
                }
        if row_data:
            sample.append({"row": row_idx, "cells": row_data})
    info["sample_data"] = sample

    return info


def validate_structure(summary):
    """Validate extracted structure against spec expectations."""
    validation = {
        "total_sheets": len(summary["sheets"]),
        "expected_sheets": 8,
        "sheets_match": len(summary["sheets"]) == 8,
        "sheet_name_mapping": {},
        "issues": [],
        "warnings": [],
    }

    actual_names = [s["sheet_name"] for s in summary["sheets"]]

    # Try to map actual sheet names to expected glossary names
    for actual_name in actual_names:
        mapped = None
        # Try common patterns
        if "目录" in actual_name or "Index" in actual_name.lower():
            mapped = "Tab_Index"
        elif "H6A" in actual_name or "程序表" in actual_name:
            mapped = "Procedure_Table_H6A"
        elif "H6-1" in actual_name or "审定" in actual_name:
            mapped = "Adjudication_H6_1"
        elif "上市" in actual_name or "Listed" in actual_name.lower():
            mapped = "Disclosure_Listed"
        elif "国有" in actual_name or "SOE" in actual_name.lower():
            mapped = "Disclosure_SOE"
        elif "H6-2" in actual_name or "明细" in actual_name:
            mapped = "Detail_H6_2"
        elif "H6-3" in actual_name or "调整" in actual_name:
            mapped = "Adjustment_H6_3"
        elif "H6-4" in actual_name or "检查" in actual_name:
            mapped = "Check_H6_4"

        validation["sheet_name_mapping"][actual_name] = mapped

    # Per-sheet validation
    for sheet_info in summary["sheets"]:
        actual_name = sheet_info["sheet_name"]
        mapped = validation["sheet_name_mapping"].get(actual_name)
        if mapped and mapped in SPEC_EXPECTED:
            expected = SPEC_EXPECTED[mapped]
            # Check row count
            if sheet_info["max_row"] != expected["rows"]:
                validation["warnings"].append(
                    f"{actual_name} (→{mapped}): 实际行数={sheet_info['max_row']}, spec预期={expected['rows']}"
                )
            # Check col count
            if sheet_info["max_column"] != expected["cols"]:
                validation["warnings"].append(
                    f"{actual_name} (→{mapped}): 实际列数={sheet_info['max_column']}, spec预期={expected['cols']}"
                )
            # Check formula count if specified
            if "formulas" in expected:
                if sheet_info["formula_count"] != expected["formulas"]:
                    validation["warnings"].append(
                        f"{actual_name} (→{mapped}): 实际公式数={sheet_info['formula_count']}, spec预期={expected['formulas']}"
                    )
        elif mapped is None:
            validation["issues"].append(f"无法映射sheet名: '{actual_name}' → spec Glossary")

    # Check all expected sheets are covered
    mapped_values = set(v for v in validation["sheet_name_mapping"].values() if v)
    missing = set(EXPECTED_SHEETS) - mapped_values
    if missing:
        validation["issues"].append(f"以下spec预期sheet未找到映射: {list(missing)}")

    validation["is_valid"] = len(validation["issues"]) == 0

    return validation


def main():
    print(f"读取: {XLSX_PATH}")
    if not XLSX_PATH.exists():
        print(f"ERROR: 文件不存在: {XLSX_PATH}")
        return

    # Load workbook with data_only=False to see formulas
    wb = load_workbook(str(XLSX_PATH), data_only=False, read_only=False)

    print(f"Sheet数量: {len(wb.sheetnames)}")
    print(f"Sheet名: {wb.sheetnames}")

    summary = {
        "source_file": str(XLSX_PATH),
        "sheet_count": len(wb.sheetnames),
        "sheet_names": wb.sheetnames,
        "sheets": [],
    }

    for sheet_name in wb.sheetnames:
        ws = wb[sheet_name]
        print(f"\n处理sheet: {sheet_name} ({ws.max_row}行×{ws.max_column}列)")
        info = extract_sheet_info(ws)
        summary["sheets"].append(info)
        print(f"  公式数: {info['formula_count']}, 合并区域: {info['merged_count']}")

    # Validate against spec
    validation = validate_structure(summary)
    summary["validation"] = validation

    print("\n" + "=" * 60)
    print("验证结果:")
    print(f"  Sheet数量匹配: {validation['sheets_match']} ({validation['total_sheets']}/8)")
    print(f"  Sheet映射: {json.dumps(validation['sheet_name_mapping'], ensure_ascii=False, indent=4)}")
    if validation["issues"]:
        print(f"  问题: {validation['issues']}")
    if validation["warnings"]:
        print(f"  警告:")
        for w in validation["warnings"]:
            print(f"    - {w}")
    print(f"  总体验证: {'✓ 通过' if validation['is_valid'] else '✗ 有问题'}")

    # Write output
    OUTPUT_PATH.parent.mkdir(parents=True, exist_ok=True)
    with open(OUTPUT_PATH, 'w', encoding='utf-8') as f:
        json.dump(summary, f, ensure_ascii=False, indent=2)

    print(f"\n产出已写入: {OUTPUT_PATH}")

    wb.close()


if __name__ == "__main__":
    main()

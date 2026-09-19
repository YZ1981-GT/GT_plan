"""
Phase 0 双源输入 - Task 0.1
openpyxl脚本读取 I4 长期待摊费用.xlsx 全部12 sheet
提取: sheet names, column headers, row counts, merged cells, formulas (data_only=False)
输出: i4_structure_summary.json
"""
import json
import os
from pathlib import Path

import openpyxl

# 源文件路径
XLSX_PATH = Path(r"d:\GT_plan\基础数据\致同通用审计程序及底稿模板（2025年修订）\1.致同审计程序及底稿模板（2025年）\4.风险应对-实质性程序（D-N）\I 无形资产循环\I4 长期待摊费用.xlsx")
OUTPUT_PATH = Path(r"d:\GT_plan\.kiro\specs\i4-long-term-prepaid\i4_structure_summary.json")


def extract_sheet_structure(ws):
    """提取单个sheet的结构信息"""
    sheet_info = {
        "title": ws.title,
        "dimensions": ws.dimensions,
        "max_row": ws.max_row,
        "max_column": ws.max_column,
        "merged_cells": [str(mc) for mc in ws.merged_cells.ranges],
        "merged_cell_count": len(ws.merged_cells.ranges),
        "column_headers": [],
        "formulas": [],
        "formula_count": 0,
        "sample_data_rows": [],
    }

    # 提取列头（前3行，通常含表头）
    for row_idx in range(1, min(4, ws.max_row + 1)):
        row_data = []
        for col_idx in range(1, ws.max_column + 1):
            cell = ws.cell(row=row_idx, column=col_idx)
            val = cell.value
            if val is not None:
                row_data.append({
                    "col": col_idx,
                    "value": str(val)[:200],  # 截断过长值
                    "is_formula": str(val).startswith("=") if isinstance(val, str) else False,
                })
        if row_data:
            sheet_info["column_headers"].append({
                "row": row_idx,
                "cells": row_data,
            })

    # 扫描所有单元格，提取公式
    formulas = []
    for row in ws.iter_rows(min_row=1, max_row=ws.max_row, max_col=ws.max_column):
        for cell in row:
            if isinstance(cell.value, str) and cell.value.startswith("="):
                formulas.append({
                    "cell": cell.coordinate,
                    "formula": cell.value[:300],  # 截断过长公式
                })

    sheet_info["formulas"] = formulas
    sheet_info["formula_count"] = len(formulas)

    # 提取若干示例数据行（第4~8行，帮助理解结构）
    for row_idx in range(4, min(9, ws.max_row + 1)):
        row_data = []
        for col_idx in range(1, ws.max_column + 1):
            cell = ws.cell(row=row_idx, column=col_idx)
            val = cell.value
            if val is not None:
                cell_info = {
                    "col": col_idx,
                    "value": str(val)[:200],
                }
                if isinstance(val, str) and val.startswith("="):
                    cell_info["is_formula"] = True
                row_data.append(cell_info)
        if row_data:
            sheet_info["sample_data_rows"].append({
                "row": row_idx,
                "cells": row_data,
            })

    return sheet_info


def main():
    print(f"读取文件: {XLSX_PATH}")
    if not XLSX_PATH.exists():
        print(f"错误: 文件不存在 - {XLSX_PATH}")
        return

    # data_only=False 以读取公式
    wb = openpyxl.load_workbook(str(XLSX_PATH), data_only=False)

    summary = {
        "source_file": str(XLSX_PATH),
        "sheet_count": len(wb.sheetnames),
        "sheet_names": wb.sheetnames,
        "sheets": [],
    }

    print(f"共 {len(wb.sheetnames)} 个sheet:")
    for i, name in enumerate(wb.sheetnames, 1):
        print(f"  [{i}] {name}")

    for sheet_name in wb.sheetnames:
        ws = wb[sheet_name]
        print(f"\n处理 sheet: {sheet_name} ({ws.max_row}行 × {ws.max_column}列)")
        sheet_info = extract_sheet_structure(ws)
        summary["sheets"].append(sheet_info)
        print(f"  合并单元格: {sheet_info['merged_cell_count']}")
        print(f"  公式数: {sheet_info['formula_count']}")

    # 生成汇总统计
    summary["statistics"] = {
        "total_formulas": sum(s["formula_count"] for s in summary["sheets"]),
        "total_merged_cells": sum(s["merged_cell_count"] for s in summary["sheets"]),
        "sheets_by_complexity": sorted(
            [{"name": s["title"], "formulas": s["formula_count"], "rows": s["max_row"], "cols": s["max_column"]}
             for s in summary["sheets"]],
            key=lambda x: x["formulas"],
            reverse=True,
        ),
    }

    # 写入JSON
    OUTPUT_PATH.parent.mkdir(parents=True, exist_ok=True)
    with open(OUTPUT_PATH, "w", encoding="utf-8") as f:
        json.dump(summary, f, ensure_ascii=False, indent=2)

    print(f"\n✅ 结构摘要已输出: {OUTPUT_PATH}")
    print(f"   Sheet总数: {summary['sheet_count']}")
    print(f"   公式总数: {summary['statistics']['total_formulas']}")
    print(f"   合并单元格总数: {summary['statistics']['total_merged_cells']}")


if __name__ == "__main__":
    main()

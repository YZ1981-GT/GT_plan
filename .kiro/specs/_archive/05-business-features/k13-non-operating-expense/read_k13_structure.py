"""
Phase0 Task 0.1: 读取 K13 营业外支出.xlsx 全部 sheet 结构
产出：k13_structure_summary.json（确认 K13-1 ~69公式 / K13-2 ~31公式）
"""

import json
import os
from pathlib import Path
from openpyxl import load_workbook
from openpyxl.utils import get_column_letter

XLSX_PATH = (
    r"D:\GT_plan\基础数据\致同通用审计程序及底稿模板（2025年修订）"
    r"\1.致同审计程序及底稿模板（2025年）"
    r"\4.风险应对-实质性程序（D-N）\K 管理循环\K13 营业外支出.xlsx"
)

OUTPUT_PATH = Path(r"d:\GT_plan\.kiro\specs\k13-non-operating-expense\k13_structure_summary.json")


def analyze_sheet(ws):
    """分析单个 sheet 结构"""
    info = {
        "title": ws.title,
        "dimensions": ws.dimensions,
        "max_row": ws.max_row,
        "max_col": ws.max_column,
        "merged_cells": [str(mc) for mc in ws.merged_cells.ranges],
        "headers": [],
        "formula_count": 0,
        "formulas": [],
        "data_sample": [],
    }

    # 提取前3行作为表头候选
    for row_idx in range(1, min(4, ws.max_row + 1)):
        row_data = []
        for col_idx in range(1, ws.max_column + 1):
            cell = ws.cell(row=row_idx, column=col_idx)
            val = cell.value
            if val is not None:
                row_data.append(f"{get_column_letter(col_idx)}{row_idx}: {val}")
        if row_data:
            info["headers"].append(row_data)

    # 扫描全部单元格计算公式数量
    for row in ws.iter_rows(min_row=1, max_row=ws.max_row,
                            min_col=1, max_col=ws.max_column):
        for cell in row:
            if isinstance(cell.value, str) and cell.value.startswith("="):
                info["formula_count"] += 1
                # 记录前20个公式的位置和内容
                if len(info["formulas"]) < 20:
                    info["formulas"].append({
                        "cell": f"{get_column_letter(cell.column)}{cell.row}",
                        "formula": cell.value
                    })

    # 提取数据样本（第4~8行）
    for row_idx in range(4, min(9, ws.max_row + 1)):
        row_data = {}
        for col_idx in range(1, ws.max_column + 1):
            cell = ws.cell(row=row_idx, column=col_idx)
            if cell.value is not None:
                col_letter = get_column_letter(col_idx)
                row_data[f"{col_letter}{row_idx}"] = str(cell.value)
        if row_data:
            info["data_sample"].append(row_data)

    return info


def main():
    print(f"读取工作簿: {XLSX_PATH}")
    assert os.path.exists(XLSX_PATH), f"文件不存在: {XLSX_PATH}"

    wb = load_workbook(XLSX_PATH, data_only=False)  # data_only=False 保留公式

    result = {
        "file": XLSX_PATH,
        "sheet_count": len(wb.sheetnames),
        "sheet_names": wb.sheetnames,
        "sheets": {},
        "summary": {},
    }

    total_formulas = 0
    for sheet_name in wb.sheetnames:
        ws = wb[sheet_name]
        sheet_info = analyze_sheet(ws)
        result["sheets"][sheet_name] = sheet_info
        total_formulas += sheet_info["formula_count"]
        print(f"  Sheet [{sheet_name}]: {ws.max_row}行 x {ws.max_column}列, "
              f"公式={sheet_info['formula_count']}, "
              f"合并单元格={len(sheet_info['merged_cells'])}")

    result["summary"] = {
        "total_formulas": total_formulas,
        "formula_by_sheet": {
            name: result["sheets"][name]["formula_count"]
            for name in wb.sheetnames
        },
        "confirmation": {
            "K13-1_formula_count": None,
            "K13-2_formula_count": None,
            "K13-1_target": 69,
            "K13-2_target": 31,
        }
    }

    # 查找 K13-1 和 K13-2 的公式数
    for name in wb.sheetnames:
        if "K13-1" in name or "审定" in name:
            result["summary"]["confirmation"]["K13-1_formula_count"] = result["sheets"][name]["formula_count"]
            result["summary"]["confirmation"]["K13-1_sheet_name"] = name
        if "K13-2" in name or "明细" in name:
            result["summary"]["confirmation"]["K13-2_formula_count"] = result["sheets"][name]["formula_count"]
            result["summary"]["confirmation"]["K13-2_sheet_name"] = name

    # 验证
    k13_1_count = result["summary"]["confirmation"]["K13-1_formula_count"]
    k13_2_count = result["summary"]["confirmation"]["K13-2_formula_count"]
    print(f"\n=== 验证 ===")
    print(f"K13-1 公式数: {k13_1_count} (目标 ~69)")
    print(f"K13-2 公式数: {k13_2_count} (目标 ~31)")
    print(f"总公式数: {total_formulas}")

    if k13_1_count:
        result["summary"]["confirmation"]["K13-1_match"] = abs(k13_1_count - 69) <= 5
    if k13_2_count:
        result["summary"]["confirmation"]["K13-2_match"] = abs(k13_2_count - 31) <= 5

    # 输出 JSON
    OUTPUT_PATH.parent.mkdir(parents=True, exist_ok=True)
    with open(OUTPUT_PATH, "w", encoding="utf-8") as f:
        json.dump(result, f, ensure_ascii=False, indent=2)

    print(f"\n产出文件: {OUTPUT_PATH}")
    print("Done!")


if __name__ == "__main__":
    main()

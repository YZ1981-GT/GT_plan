"""
Phase 0.1: openpyxl脚本读取L3长期借款.xlsx全部14 sheet
提取结构 + 确认利息测算表L3-5+明细表L3-2（一年内到期列）
产出：l3_structure_summary.json
"""

import json
import os
from pathlib import Path

import openpyxl
from openpyxl.utils import get_column_letter

# 源文件路径
XLSX_PATH = Path(r"d:\GT_plan\数据\致同通用审计程序及底稿模板（2025年修订）\1.致同审计程序及底稿模板（2025年）\4.风险应对-实质性程序（D-N）\L 债务循环\L3 长期借款.xlsx")
OUTPUT_PATH = Path(r"d:\GT_plan\.kiro\specs\l3-long-term-loans\l3_structure_summary.json")


def extract_sheet_structure(ws):
    """提取单个sheet的结构信息"""
    info = {
        "title": ws.title,
        "dimensions": ws.dimensions,
        "max_row": ws.max_row,
        "max_column": ws.max_column,
        "row_count": ws.max_row,
        "col_count": ws.max_column,
        "merged_cells": [str(m) for m in ws.merged_cells.ranges],
        "headers": [],
        "formulas": [],
        "sample_data": [],
    }

    # 提取前5行作为header候选
    for row_idx in range(1, min(6, ws.max_row + 1)):
        row_data = []
        for col_idx in range(1, ws.max_column + 1):
            cell = ws.cell(row=row_idx, column=col_idx)
            val = cell.value
            if val is not None:
                row_data.append({
                    "col": get_column_letter(col_idx),
                    "col_idx": col_idx,
                    "value": str(val)[:100],
                    "is_formula": str(val).startswith("=") if isinstance(val, str) else False,
                })
        if row_data:
            info["headers"].append({"row": row_idx, "cells": row_data})

    # 扫描所有公式
    formula_count = 0
    formula_samples = []
    for row in ws.iter_rows(min_row=1, max_row=ws.max_row, max_col=ws.max_column):
        for cell in row:
            if isinstance(cell.value, str) and cell.value.startswith("="):
                formula_count += 1
                if len(formula_samples) < 15:
                    formula_samples.append({
                        "cell": f"{get_column_letter(cell.column)}{cell.row}",
                        "formula": cell.value[:200],
                    })

    info["formula_count"] = formula_count
    info["formula_samples"] = formula_samples

    # 提取第6~10行sample data (如有)
    for row_idx in range(6, min(11, ws.max_row + 1)):
        row_data = []
        for col_idx in range(1, min(ws.max_column + 1, 20)):
            cell = ws.cell(row=row_idx, column=col_idx)
            val = cell.value
            if val is not None:
                row_data.append({
                    "col": get_column_letter(col_idx),
                    "value": str(val)[:100],
                    "is_formula": str(val).startswith("=") if isinstance(val, str) else False,
                })
        if row_data:
            info["sample_data"].append({"row": row_idx, "cells": row_data})

    return info


def confirm_l3_5_interest(wb):
    """确认利息测算表L3-5的结构"""
    confirmation = {"found": False, "sheet_name": None, "details": {}}

    for name in wb.sheetnames:
        if "L3-5" in name or "利息" in name:
            ws = wb[name]
            confirmation["found"] = True
            confirmation["sheet_name"] = name
            confirmation["details"] = {
                "rows": ws.max_row,
                "cols": ws.max_column,
                "headers_row1": [],
                "headers_row2": [],
                "has_principal_col": False,
                "has_rate_col": False,
                "has_days_col": False,
                "has_interest_calc_col": False,
            }

            # 检查前3行header
            for row_idx in range(1, min(4, ws.max_row + 1)):
                for col_idx in range(1, ws.max_column + 1):
                    val = ws.cell(row=row_idx, column=col_idx).value
                    if val:
                        val_str = str(val)
                        key = f"headers_row{row_idx}"
                        if key in confirmation["details"]:
                            confirmation["details"][key].append({
                                "col": get_column_letter(col_idx),
                                "value": val_str[:80]
                            })
                        # 检查关键列
                        if "本金" in val_str or "金额" in val_str:
                            confirmation["details"]["has_principal_col"] = True
                        if "利率" in val_str:
                            confirmation["details"]["has_rate_col"] = True
                        if "天数" in val_str or "期间" in val_str:
                            confirmation["details"]["has_days_col"] = True
                        if "测算" in val_str or "计算" in val_str:
                            confirmation["details"]["has_interest_calc_col"] = True
            break

    return confirmation


def confirm_l3_2_detail(wb):
    """确认明细表L3-2的结构（一年内到期列）"""
    confirmation = {"found": False, "sheet_name": None, "details": {}}

    for name in wb.sheetnames:
        if "L3-2" in name or "明细" in name:
            ws = wb[name]
            confirmation["found"] = True
            confirmation["sheet_name"] = name
            confirmation["details"] = {
                "rows": ws.max_row,
                "cols": ws.max_column,
                "all_headers": [],
                "has_due_date_col": False,
                "has_current_portion_col": False,
                "has_loan_type_col": False,
                "has_rate_col": False,
                "has_begin_balance_col": False,
                "has_end_balance_col": False,
            }

            # 检查前4行header
            for row_idx in range(1, min(5, ws.max_row + 1)):
                for col_idx in range(1, ws.max_column + 1):
                    val = ws.cell(row=row_idx, column=col_idx).value
                    if val:
                        val_str = str(val)
                        confirmation["details"]["all_headers"].append({
                            "row": row_idx,
                            "col": get_column_letter(col_idx),
                            "col_idx": col_idx,
                            "value": val_str[:80]
                        })
                        # 检查关键列
                        if "到期" in val_str and "日" in val_str:
                            confirmation["details"]["has_due_date_col"] = True
                        if "一年内" in val_str or "流动" in val_str:
                            confirmation["details"]["has_current_portion_col"] = True
                        if "类型" in val_str or "种类" in val_str:
                            confirmation["details"]["has_loan_type_col"] = True
                        if "利率" in val_str:
                            confirmation["details"]["has_rate_col"] = True
                        if "期初" in val_str:
                            confirmation["details"]["has_begin_balance_col"] = True
                        if "期末" in val_str:
                            confirmation["details"]["has_end_balance_col"] = True
            break

    return confirmation


def main():
    if not XLSX_PATH.exists():
        print(f"ERROR: 源文件不存在: {XLSX_PATH}")
        return

    print(f"正在读取: {XLSX_PATH}")
    wb = openpyxl.load_workbook(str(XLSX_PATH), data_only=False)

    print(f"总sheet数: {len(wb.sheetnames)}")
    print(f"Sheet列表: {wb.sheetnames}")

    result = {
        "source_file": str(XLSX_PATH),
        "total_sheets": len(wb.sheetnames),
        "sheet_names": wb.sheetnames,
        "sheets": {},
        "confirmation": {
            "l3_5_interest_calc": None,
            "l3_2_detail": None,
        },
        "summary": {
            "total_formulas": 0,
            "sheets_with_formulas": [],
            "key_findings": [],
        }
    }

    # 逐sheet提取结构
    for name in wb.sheetnames:
        print(f"  处理: {name}")
        ws = wb[name]
        info = extract_sheet_structure(ws)
        result["sheets"][name] = info

        if info["formula_count"] > 0:
            result["summary"]["total_formulas"] += info["formula_count"]
            result["summary"]["sheets_with_formulas"].append({
                "name": name,
                "formula_count": info["formula_count"]
            })

    # 确认L3-5利息测算表
    result["confirmation"]["l3_5_interest_calc"] = confirm_l3_5_interest(wb)

    # 确认L3-2明细表（一年内到期列）
    result["confirmation"]["l3_2_detail"] = confirm_l3_2_detail(wb)

    # 生成key findings
    findings = result["summary"]["key_findings"]
    if result["confirmation"]["l3_5_interest_calc"]["found"]:
        det = result["confirmation"]["l3_5_interest_calc"]["details"]
        findings.append(f"✅ L3-5利息测算表已确认: {result['confirmation']['l3_5_interest_calc']['sheet_name']} ({det['rows']}行×{det['cols']}列)")
        if det.get("has_principal_col"):
            findings.append("  ✅ 含本金列")
        if det.get("has_rate_col"):
            findings.append("  ✅ 含利率列")
        if det.get("has_days_col"):
            findings.append("  ✅ 含天数/期间列")
        if det.get("has_interest_calc_col"):
            findings.append("  ✅ 含测算利息列")
    else:
        findings.append("⚠️ 未找到L3-5利息测算表")

    if result["confirmation"]["l3_2_detail"]["found"]:
        det = result["confirmation"]["l3_2_detail"]["details"]
        findings.append(f"✅ L3-2明细表已确认: {result['confirmation']['l3_2_detail']['sheet_name']} ({det['rows']}行×{det['cols']}列)")
        if det.get("has_current_portion_col"):
            findings.append("  ✅ 含一年内到期列")
        else:
            findings.append("  ⚠️ 未直接找到'一年内到期'列名（可能在公式或合并单元格中）")
        if det.get("has_due_date_col"):
            findings.append("  ✅ 含到期日列")
        if det.get("has_rate_col"):
            findings.append("  ✅ 含利率列")
    else:
        findings.append("⚠️ 未找到L3-2明细表")

    findings.append(f"📊 总公式数: {result['summary']['total_formulas']}")

    # 输出JSON
    OUTPUT_PATH.parent.mkdir(parents=True, exist_ok=True)
    with open(OUTPUT_PATH, "w", encoding="utf-8") as f:
        json.dump(result, f, ensure_ascii=False, indent=2)

    print(f"\n产出文件: {OUTPUT_PATH}")
    print(f"\n=== Key Findings ===")
    for finding in findings:
        print(f"  {finding}")

    wb.close()


if __name__ == "__main__":
    main()

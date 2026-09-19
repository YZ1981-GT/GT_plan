"""
Phase 0.1: openpyxl脚本读取N5所得税费用.xlsx全部15 sheet
提取: sheet名/列头/行数/公式单元格/合并区域/数据类型
标记skip: N3A原底稿
产出: n5_structure_summary.json（权威列头+公式清单）
"""

import json
import os
from openpyxl import load_workbook
from openpyxl.utils import get_column_letter

# Source xlsx path
XLSX_PATH = (
    r"D:\GT_plan\数据\致同通用审计程序及底稿模板（2025年修订）"
    r"\1.致同审计程序及底稿模板（2025年）\4.风险应对-实质性程序（D-N）"
    r"\N 税金循环\N5 所得税费用.xlsx"
)

# Output path
OUTPUT_DIR = os.path.dirname(os.path.abspath(__file__))
OUTPUT_PATH = os.path.join(OUTPUT_DIR, "n5_structure_summary.json")

# Sheets to skip (辅助sheet, 走OnlyOffice fallback)
SKIP_SHEETS = ["N3A"]


def is_skip_sheet(sheet_name: str) -> bool:
    """Check if sheet should be marked as skip (N3A原底稿)."""
    for skip in SKIP_SHEETS:
        if skip in sheet_name:
            return True
    return False


def extract_column_headers(ws, max_header_rows=3):
    """Extract column headers from first few rows."""
    headers = []
    for row_idx in range(1, min(max_header_rows + 1, ws.max_row + 1)):
        row_headers = []
        for col_idx in range(1, ws.max_column + 1):
            cell = ws.cell(row=row_idx, column=col_idx)
            val = cell.value
            if val is not None:
                row_headers.append({
                    "col": get_column_letter(col_idx),
                    "col_idx": col_idx,
                    "value": str(val).strip()
                })
        if row_headers:
            headers.append({"row": row_idx, "cells": row_headers})
    return headers


def extract_formulas(ws):
    """Extract all formula cells."""
    formulas = []
    for row in ws.iter_rows(min_row=1, max_row=ws.max_row,
                            min_col=1, max_col=ws.max_column):
        for cell in row:
            if cell.value and isinstance(cell.value, str) and cell.value.startswith("="):
                formulas.append({
                    "cell": f"{get_column_letter(cell.column)}{cell.row}",
                    "row": cell.row,
                    "col": cell.column,
                    "formula": cell.value
                })
    return formulas


def extract_data_types(ws, sample_rows=10):
    """Extract data type distribution from sample rows."""
    type_map = {}
    start_row = min(4, ws.max_row)  # Skip header rows
    end_row = min(start_row + sample_rows, ws.max_row + 1)

    for row in ws.iter_rows(min_row=start_row, max_row=end_row,
                            min_col=1, max_col=ws.max_column):
        for cell in row:
            col_letter = get_column_letter(cell.column)
            if col_letter not in type_map:
                type_map[col_letter] = set()
            if cell.value is not None:
                if isinstance(cell.value, str) and cell.value.startswith("="):
                    type_map[col_letter].add("formula")
                elif isinstance(cell.value, (int, float)):
                    type_map[col_letter].add("numeric")
                elif isinstance(cell.value, str):
                    type_map[col_letter].add("string")
                else:
                    type_map[col_letter].add(type(cell.value).__name__)

    # Convert sets to lists for JSON serialization
    return {k: list(v) for k, v in type_map.items()}


def extract_merged_regions(ws):
    """Extract merged cell regions."""
    return [str(m) for m in ws.merged_cells.ranges]


def analyze_sheet(ws, sheet_name):
    """Analyze a single sheet."""
    skip = is_skip_sheet(sheet_name)

    result = {
        "sheet_name": sheet_name,
        "skip": skip,
        "dimensions": {
            "rows": ws.max_row,
            "cols": ws.max_column,
            "size": f"{ws.max_row}×{ws.max_column}"
        },
    }

    if skip:
        result["skip_reason"] = "N3A原底稿，走OnlyOffice fallback，不做HTML组件化"
        return result

    # Full analysis for non-skip sheets
    headers = extract_column_headers(ws)
    formulas = extract_formulas(ws)
    merged = extract_merged_regions(ws)
    data_types = extract_data_types(ws)

    result["column_headers"] = headers
    result["formulas"] = {
        "count": len(formulas),
        "cells": formulas
    }
    result["merged_regions"] = {
        "count": len(merged),
        "ranges": merged
    }
    result["data_types"] = data_types

    return result


def main():
    print(f"Loading workbook: {XLSX_PATH}")
    if not os.path.exists(XLSX_PATH):
        print(f"ERROR: File not found: {XLSX_PATH}")
        return

    # Load with data_only=False to preserve formulas
    wb = load_workbook(XLSX_PATH, data_only=False, read_only=False)

    print(f"Total sheets: {len(wb.sheetnames)}")
    print(f"Sheet names: {wb.sheetnames}")

    summary = {
        "source_file": os.path.basename(XLSX_PATH),
        "total_sheets": len(wb.sheetnames),
        "sheet_names": wb.sheetnames,
        "skip_sheets": [],
        "active_sheets": [],
        "sheets": []
    }

    for sheet_name in wb.sheetnames:
        print(f"\nAnalyzing: {sheet_name}")
        ws = wb[sheet_name]
        sheet_data = analyze_sheet(ws, sheet_name)
        summary["sheets"].append(sheet_data)

        if sheet_data.get("skip"):
            summary["skip_sheets"].append(sheet_name)
            print(f"  → SKIP (N3A原底稿)")
        else:
            summary["active_sheets"].append(sheet_name)
            dims = sheet_data["dimensions"]
            formula_count = sheet_data["formulas"]["count"]
            merged_count = sheet_data["merged_regions"]["count"]
            print(f"  → {dims['size']}, {formula_count} formulas, {merged_count} merged regions")

    # Print validation summary
    print("\n" + "=" * 60)
    print("VALIDATION SUMMARY")
    print("=" * 60)
    print(f"Total sheets: {summary['total_sheets']}")
    print(f"Active sheets: {len(summary['active_sheets'])}")
    print(f"Skip sheets: {len(summary['skip_sheets'])} → {summary['skip_sheets']}")
    print()

    # Check expected dimensions from task description
    expected = {
        "N5-1": {"rows": 28, "cols": 14, "formulas": 36},
        "N5-2": {"rows": 38, "cols": 10, "formulas": 8},
        "N5-4": {"rows": 82, "cols": 7},
        "N5-5": {"rows": 107, "cols": 8},
        "N5-6": {"rows": 54, "cols": 6, "formulas": 10},
        "N5-6-1": {"rows": 43, "cols": 7, "formulas": 17},
        "N5-6-2": {"rows": 18, "cols": 13},
        "N5-7": {"rows": 15, "cols": 7, "formulas": 12},
        "N5-8": {"rows": 44, "cols": 10, "formulas": 12},
    }

    print("\nDimension Check (expected vs actual):")
    for sheet_data in summary["sheets"]:
        name = sheet_data["sheet_name"]
        if sheet_data.get("skip"):
            continue
        dims = sheet_data["dimensions"]
        formula_count = sheet_data["formulas"]["count"]
        # Find matching expected key
        for key, exp in expected.items():
            if key in name:
                row_match = "✓" if dims["rows"] == exp["rows"] else f"✗({dims['rows']}≠{exp['rows']})"
                col_match = "✓" if dims["cols"] == exp["cols"] else f"✗({dims['cols']}≠{exp['cols']})"
                formula_match = ""
                if "formulas" in exp:
                    formula_match = "✓" if formula_count == exp["formulas"] else f"✗({formula_count}≠{exp['formulas']})"
                print(f"  {name}: rows={row_match}, cols={col_match}, formulas={formula_match}")
                break

    # Save output
    with open(OUTPUT_PATH, "w", encoding="utf-8") as f:
        json.dump(summary, f, ensure_ascii=False, indent=2)

    print(f"\nOutput saved to: {OUTPUT_PATH}")
    wb.close()


if __name__ == "__main__":
    main()

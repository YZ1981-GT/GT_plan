"""
Phase 0.1: openpyxl脚本读取L2应付利息.xlsx全部8 sheet
提取结构 + 确认审定表L2-1（79公式）+ 明细表L2-2结构
产出：l2_structure_summary.json
"""

import json
import os
from pathlib import Path
from openpyxl import load_workbook
from openpyxl.utils import get_column_letter

# 源模板路径（使用wp_templates下的副本）
XLSX_PATH = Path(r"d:\GT_plan\backend\wp_templates\L\L2 应付利息.xlsx")
OUTPUT_PATH = Path(r"d:\GT_plan\.kiro\specs\l2-interest-payable\l2_structure_summary.json")


def extract_sheet_structure(ws):
    """提取单个sheet的结构信息"""
    info = {
        "title": ws.title,
        "dimensions": ws.dimensions,
        "max_row": ws.max_row,
        "max_column": ws.max_column,
        "size": f"{ws.max_row}×{ws.max_column}",
    }

    # 提取列头（前3行，通常包含表头）
    headers = []
    for row_idx in range(1, min(4, ws.max_row + 1)):
        row_headers = []
        for col_idx in range(1, ws.max_column + 1):
            cell = ws.cell(row=row_idx, column=col_idx)
            val = cell.value
            if val is not None:
                row_headers.append({
                    "col": get_column_letter(col_idx),
                    "col_idx": col_idx,
                    "value": str(val)[:100]  # 截断过长值
                })
        if row_headers:
            headers.append({"row": row_idx, "cells": row_headers})
    info["headers"] = headers

    # 提取公式单元格
    formulas = []
    for row in ws.iter_rows(min_row=1, max_row=ws.max_row, max_col=ws.max_column):
        for cell in row:
            if cell.value and isinstance(cell.value, str) and cell.value.startswith("="):
                formulas.append({
                    "cell": f"{get_column_letter(cell.column)}{cell.row}",
                    "formula": cell.value[:200]  # 截断过长公式
                })
    info["formula_count"] = len(formulas)
    info["formulas"] = formulas

    # 提取合并单元格
    merged = [str(m) for m in ws.merged_cells.ranges]
    info["merged_cells"] = merged
    info["merged_count"] = len(merged)

    return info


def main():
    if not XLSX_PATH.exists():
        print(f"ERROR: 文件不存在: {XLSX_PATH}")
        return

    print(f"读取文件: {XLSX_PATH}")
    # data_only=False 以保留公式
    wb = load_workbook(str(XLSX_PATH), data_only=False, read_only=False)

    result = {
        "source_file": str(XLSX_PATH),
        "sheet_count": len(wb.sheetnames),
        "sheet_names": wb.sheetnames,
        "sheets": {},
        "summary": {}
    }

    total_formulas = 0
    for sheet_name in wb.sheetnames:
        ws = wb[sheet_name]
        print(f"  处理sheet: {sheet_name} ({ws.max_row}×{ws.max_column})")
        sheet_info = extract_sheet_structure(ws)
        result["sheets"][sheet_name] = sheet_info
        total_formulas += sheet_info["formula_count"]

    # 汇总
    result["summary"] = {
        "total_sheets": len(wb.sheetnames),
        "total_formulas": total_formulas,
        "sheet_sizes": {name: result["sheets"][name]["size"] for name in wb.sheetnames},
        "formula_distribution": {name: result["sheets"][name]["formula_count"] for name in wb.sheetnames},
    }

    # 特别标注审定表L2-1和明细表L2-2
    for name, info in result["sheets"].items():
        if "审定" in name or "L2-1" in name:
            result["summary"]["adjudication_L2_1"] = {
                "sheet_name": name,
                "size": info["size"],
                "formula_count": info["formula_count"],
                "note": "审定表，预期~79公式"
            }
        if "明细" in name or "L2-2" in name:
            result["summary"]["detail_L2_2"] = {
                "sheet_name": name,
                "size": info["size"],
                "formula_count": info["formula_count"],
                "max_column": info["max_column"],
                "note": "明细表，预期27列"
            }

    wb.close()

    # 写出JSON
    OUTPUT_PATH.parent.mkdir(parents=True, exist_ok=True)
    with open(OUTPUT_PATH, "w", encoding="utf-8") as f:
        json.dump(result, f, ensure_ascii=False, indent=2)

    print(f"\n产出文件: {OUTPUT_PATH}")
    print(f"总sheet数: {result['summary']['total_sheets']}")
    print(f"总公式数: {total_formulas}")
    print(f"各sheet公式分布: {result['summary']['formula_distribution']}")

    # 验证关键预期
    adj = result["summary"].get("adjudication_L2_1", {})
    det = result["summary"].get("detail_L2_2", {})
    print(f"\n=== 关键验证 ===")
    print(f"审定表L2-1: {adj.get('sheet_name', 'NOT FOUND')} - 公式数={adj.get('formula_count', 'N/A')} (预期~79)")
    print(f"明细表L2-2: {det.get('sheet_name', 'NOT FOUND')} - 列数={det.get('max_column', 'N/A')} (预期27)")


if __name__ == "__main__":
    main()

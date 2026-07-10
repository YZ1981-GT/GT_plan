"""
Task 0.1: openpyxl脚本读取H5油气资产.xlsx全部24 sheet
提取：sheet名/列头/行数/公式单元格/合并区域/数据类型
产出：h5_structure_summary.json
"""
import json
import os
from pathlib import Path

from openpyxl import load_workbook
from openpyxl.utils import get_column_letter

# Source xlsx path
XLSX_PATH = (
    r"d:\GT_plan\基础数据\致同通用审计程序及底稿模板（2025年修订）"
    r"\1.致同审计程序及底稿模板（2025年）"
    r"\4.风险应对-实质性程序（D-N）"
    r"\H 固定资产循环"
    r"\H5 油气资产.xlsx"
)

# Output path
OUTPUT_DIR = Path(__file__).parent
OUTPUT_PATH = OUTPUT_DIR / "h5_structure_summary.json"


def extract_sheet_structure(ws):
    """Extract structure info from a single worksheet."""
    sheet_info = {
        "sheet_name": ws.title,
        "max_row": ws.max_row,
        "max_column": ws.max_column,
        "column_count": ws.max_column,
        "row_count": ws.max_row,
        "column_headers": [],
        "formula_cells": [],
        "merged_regions": [],
        "data_types": {},
    }

    # Extract column headers (first few rows that typically contain headers)
    # Try rows 1-5 for headers
    headers_by_row = {}
    for row_idx in range(1, min(6, ws.max_row + 1)):
        row_headers = []
        for col_idx in range(1, ws.max_column + 1):
            cell = ws.cell(row=row_idx, column=col_idx)
            val = cell.value
            if val is not None:
                row_headers.append({
                    "col": get_column_letter(col_idx),
                    "col_index": col_idx,
                    "value": str(val)[:200]  # Truncate long values
                })
        if row_headers:
            headers_by_row[f"row_{row_idx}"] = row_headers
    sheet_info["column_headers"] = headers_by_row

    # Extract formula cells
    formula_cells = []
    for row in ws.iter_rows(min_row=1, max_row=ws.max_row, max_col=ws.max_column):
        for cell in row:
            if cell.data_type == 'f':  # formula
                formula_cells.append({
                    "cell": f"{get_column_letter(cell.column)}{cell.row}",
                    "formula": str(cell.value)
                })
    sheet_info["formula_cells"] = formula_cells
    sheet_info["formula_count"] = len(formula_cells)

    # Extract merged regions
    merged = []
    for merged_range in ws.merged_cells.ranges:
        merged.append(str(merged_range))
    sheet_info["merged_regions"] = merged
    sheet_info["merged_count"] = len(merged)

    # Extract data types summary
    type_counts = {}
    for row in ws.iter_rows(min_row=1, max_row=ws.max_row, max_col=ws.max_column):
        for cell in row:
            dt = cell.data_type
            type_label = {
                's': 'string',
                'n': 'numeric',
                'f': 'formula',
                'b': 'boolean',
                'd': 'date',
                'e': 'error',
                None: 'empty',
            }.get(dt, dt or 'empty')
            type_counts[type_label] = type_counts.get(type_label, 0) + 1
    sheet_info["data_types"] = type_counts

    return sheet_info


def main():
    print(f"Loading workbook: {XLSX_PATH}")
    if not os.path.exists(XLSX_PATH):
        raise FileNotFoundError(f"Excel file not found: {XLSX_PATH}")

    # Load with data_only=False to see formulas
    wb = load_workbook(XLSX_PATH, data_only=False, read_only=False)

    print(f"Sheet count: {len(wb.sheetnames)}")
    print(f"Sheet names: {wb.sheetnames}")

    summary = {
        "source_file": XLSX_PATH,
        "total_sheets": len(wb.sheetnames),
        "sheet_names": wb.sheetnames,
        "sheets": []
    }

    for sheet_name in wb.sheetnames:
        print(f"  Processing: {sheet_name}")
        ws = wb[sheet_name]
        sheet_info = extract_sheet_structure(ws)
        summary["sheets"].append(sheet_info)

    # Validation: check if we have 24 sheets as expected
    summary["validation"] = {
        "expected_sheet_count": 24,
        "actual_sheet_count": len(wb.sheetnames),
        "matches_expected": len(wb.sheetnames) == 24,
    }

    # Cross-reference with spec glossary expected sheets
    expected_sheets_from_spec = [
        "Tab_Index",           # 底稿目录
        "H5A",                 # 程序表
        "H5-1",                # 审定表
        "附注_上市",           # 附注上市公司
        "附注_国企",           # 附注国企
        "H5-2",                # 明细表
        "H5-3",                # 调整分录
        "H5-4",                # 闲置检查
        "H5-5",                # 会计政策
        "H5-6",                # 分析表
        "H5-7",                # 增加检查
        "H5-8",                # 减少检查
        "H5-9",                # 监盘计划
        "H5-10",               # 盘点检查
        "H5-11",               # 监盘小结
        "H5-12(不含减值)",     # 折耗不含减值
        "H5-12(含减值)",       # 折耗含减值
        "H5-13",               # 折耗分配
        "H5-14",               # 减值测算
        "H5-15",               # 可收回金额
        "H5-16",               # 权属检查
        "H5-17",               # 关联交易
        "H5-18",               # 经营租出
        "H5-19",               # 融资租出
    ]
    summary["validation"]["expected_sheet_names_from_spec"] = expected_sheets_from_spec

    # Write output
    with open(OUTPUT_PATH, 'w', encoding='utf-8') as f:
        json.dump(summary, f, ensure_ascii=False, indent=2)

    print(f"\nOutput written to: {OUTPUT_PATH}")
    print(f"Total sheets: {summary['total_sheets']}")
    print(f"Validation: {'PASS' if summary['validation']['matches_expected'] else 'FAIL'}")
    print(f"  Expected: {summary['validation']['expected_sheet_count']}")
    print(f"  Actual: {summary['validation']['actual_sheet_count']}")

    # Print per-sheet summary
    print("\n--- Per-sheet Summary ---")
    for s in summary["sheets"]:
        print(f"  {s['sheet_name']:30s} | rows={s['row_count']:4d} | cols={s['column_count']:3d} | formulas={s['formula_count']:4d} | merged={s['merged_count']:3d}")


if __name__ == "__main__":
    main()

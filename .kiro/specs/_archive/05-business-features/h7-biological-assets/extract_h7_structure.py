"""
Phase 0 Task 0.1: openpyxl脚本读取H7生产性生物资产.xlsx全部26 sheet
提取：sheet名/列头/行数/公式/合并区域/数据类型
产出：h7_structure_summary.json
验证：26 sheet结构（含双模式sheet对）与本spec一致
"""
import json
import re
import sys
from datetime import datetime, date, time
from pathlib import Path

from openpyxl import load_workbook
from openpyxl.utils import get_column_letter


class DateTimeEncoder(json.JSONEncoder):
    """处理datetime等非JSON序列化类型"""
    def default(self, obj):
        if isinstance(obj, (datetime, date)):
            return obj.isoformat()
        if isinstance(obj, time):
            return obj.isoformat()
        return super().default(obj)

# --- 路径配置 ---
PROJECT_ROOT = Path(__file__).resolve().parents[3]  # d:\GT_plan
XLSX_PATH = PROJECT_ROOT / "backend" / "wp_templates" / "H" / "H7 生产性生物资产.xlsx"
OUTPUT_PATH = Path(__file__).resolve().parent / "h7_structure_summary.json"

# 备用路径
ALT_XLSX_PATH = (
    PROJECT_ROOT
    / "基础数据"
    / "致同通用审计程序及底稿模板（2025年修订）"
    / "1.致同审计程序及底稿模板（2025年）"
    / "4.风险应对-实质性程序（D-N）"
    / "H 固定资产循环"
    / "H7 生产性生物资产.xlsx"
)

# 双计量模式sheet对识别
DUAL_MODE_PATTERNS = {
    "H7-1": {"cost": None, "fair": None},
    "H7-2": {"cost": None, "fair": None},
    "H7-6": {"cost": None, "fair": None},
    "H7-7": {"cost": None, "fair": None},
}


def anonymize_formula(formula: str) -> str:
    """将公式中的具体单元格引用替换为REF，提取公式模式"""
    if not formula:
        return formula
    # 替换跨sheet引用 'SheetName'!A1 → 'REF'!REF
    pattern = re.sub(r"'[^']+'![A-Z]+\d+", "'REF'!REF", formula)
    # 替换单元格引用 A1, AA123 → REF
    pattern = re.sub(r"\b[A-Z]{1,3}\d+\b", "REF", pattern)
    # 替换范围 REF:REF 保持不变
    return pattern


def is_header_ref_formula(formula: str) -> bool:
    """判断公式是否为表头引用（如=底稿目录!A2）"""
    return bool(re.match(r"^=底稿目录![A-Z]+\d+$", formula))


def extract_sheet_info(ws):
    """提取单个sheet的结构信息"""
    sheet_info = {
        "title": ws.title,
        "dimensions": ws.dimensions,
        "max_row": ws.max_row,
        "max_column": ws.max_column,
        "size": f"{ws.max_row}×{ws.max_column}",
        "merged_cells": [str(m) for m in ws.merged_cells.ranges],
        "merged_count": len(ws.merged_cells.ranges),
    }

    # 提取列头（前5行）
    headers = {}
    for row_idx in range(1, min(6, ws.max_row + 1)):
        row_data = []
        for col_idx in range(1, ws.max_column + 1):
            cell = ws.cell(row=row_idx, column=col_idx)
            if cell.value is not None:
                col_letter = get_column_letter(col_idx)
                val = cell.value
                if isinstance(val, str) and val.startswith("="):
                    row_data.append({"col": col_letter, "value": val})
                elif isinstance(val, (datetime, date, time)):
                    row_data.append({"col": col_letter, "value": val.isoformat()})
                else:
                    row_data.append({"col": col_letter, "value": str(val) if not isinstance(val, (int, float)) else val})
        if row_data:
            headers[f"row{row_idx}"] = row_data
    sheet_info["headers"] = headers

    # 提取公式
    formula_cells = []
    business_formula_cells = []
    header_ref_count = 0
    formula_patterns_set = set()

    for row in ws.iter_rows(min_row=1, max_row=ws.max_row, max_col=ws.max_column):
        for cell in row:
            if cell.value and isinstance(cell.value, str) and cell.value.startswith("="):
                cell_ref = f"{get_column_letter(cell.column)}{cell.row}"
                formula_entry = {"cell": cell_ref, "formula": cell.value}
                formula_cells.append(formula_entry)

                if is_header_ref_formula(cell.value):
                    header_ref_count += 1
                else:
                    business_formula_cells.append(formula_entry)
                    pattern = anonymize_formula(cell.value)
                    formula_patterns_set.add(pattern)

    sheet_info["formula_cells"] = formula_cells
    sheet_info["formula_count"] = len(formula_cells)
    sheet_info["business_formula_cells"] = business_formula_cells
    sheet_info["business_formula_count"] = len(business_formula_cells)
    sheet_info["header_ref_count"] = header_ref_count
    sheet_info["formula_patterns"] = sorted(formula_patterns_set)

    # 业务公式行数（去重行号）
    biz_rows = set()
    for f in business_formula_cells:
        row_num = int(re.search(r"\d+", f["cell"]).group())
        biz_rows.add(row_num)
    sheet_info["business_formula_rows"] = len(biz_rows)

    # 提取样本数据行（前8行有数据的行）
    sample_rows = []
    count = 0
    for row_idx in range(1, ws.max_row + 1):
        if count >= 8:
            break
        row_data = {}
        for col_idx in range(1, ws.max_column + 1):
            cell = ws.cell(row=row_idx, column=col_idx)
            if cell.value is not None:
                col_letter = get_column_letter(col_idx)
                val = cell.value
                if isinstance(val, str) and val.startswith("="):
                    row_data[col_letter] = f"[FORMULA] {val}"
                elif isinstance(val, (datetime, date, time)):
                    row_data[col_letter] = val.isoformat()
                else:
                    row_data[col_letter] = val
        if row_data:
            sample_rows.append({"row": row_idx, "data": row_data})
            count += 1
    sheet_info["sample_data_rows"] = sample_rows

    # 数据类型统计
    type_counts = {"s": 0, "n": 0, "f": 0, "b": 0, "d": 0}
    for row in ws.iter_rows(min_row=1, max_row=ws.max_row, max_col=ws.max_column):
        for cell in row:
            if cell.value is None:
                continue
            if isinstance(cell.value, str):
                if cell.value.startswith("="):
                    type_counts["f"] += 1
                else:
                    type_counts["s"] += 1
            elif isinstance(cell.value, (int, float)):
                type_counts["n"] += 1
            elif isinstance(cell.value, bool):
                type_counts["b"] += 1
            else:
                type_counts["d"] += 1

    # 只保留非零的类型
    sheet_info["data_types"] = {k: v for k, v in type_counts.items() if v > 0}

    return sheet_info


def identify_dual_mode_pairs(sheet_names: list) -> dict:
    """识别双计量模式sheet对"""
    pairs = {
        "H7-1": {"cost": None, "fair": None},
        "H7-2": {"cost": None, "fair": None},
        "H7-6": {"cost": None, "fair": None},
        "H7-7": {"cost": None, "fair": None},
    }

    for name in sheet_names:
        for code in pairs:
            # 精确匹配：sheet名必须以code结尾或code后面紧跟非数字
            # 避免H7-1匹配H7-13/H7-14等
            pattern = re.compile(rf"{re.escape(code)}(?!\d)")
            if pattern.search(name):
                if "成本" in name or "原值" in name:
                    pairs[code]["cost"] = name
                elif "公允" in name:
                    pairs[code]["fair"] = name

    return pairs


def main():
    # 确定xlsx路径
    xlsx_path = XLSX_PATH if XLSX_PATH.exists() else ALT_XLSX_PATH
    if not xlsx_path.exists():
        print(f"ERROR: xlsx文件不存在: {XLSX_PATH}")
        print(f"  备用路径也不存在: {ALT_XLSX_PATH}")
        sys.exit(1)

    print(f"正在读取: {xlsx_path}")
    wb = load_workbook(str(xlsx_path), data_only=False, read_only=False)

    sheet_names = wb.sheetnames
    total_sheets = len(sheet_names)
    print(f"发现 {total_sheets} 个sheet")

    # 识别双模式对
    dual_mode_pairs = identify_dual_mode_pairs(sheet_names)

    # 识别skip sheets (GT_Custom等辅助sheet)
    skip_sheets = []
    active_sheets = []
    for name in sheet_names:
        if name == "GT_Custom" or "原底稿" in name or "示例" in name:
            skip_sheets.append(name)
        else:
            active_sheets.append(name)

    # 识别key_sheets（含公式的主要sheet）
    key_sheets = {}

    # 提取每个sheet结构
    sheets_data = []
    for name in sheet_names:
        print(f"  处理: {name}")
        ws = wb[name]

        if name in skip_sheets:
            sheets_data.append({
                "title": name,
                "skip": True,
                "dimensions": ws.dimensions,
                "max_row": ws.max_row,
                "max_column": ws.max_column,
                "size": f"{ws.max_row}×{ws.max_column}",
                "reason": "辅助sheet/参考底稿，不做HTML组件化",
            })
            continue

        info = extract_sheet_info(ws)
        sheets_data.append(info)

        # 如果有业务公式，加入key_sheets
        if info["business_formula_count"] > 0:
            # 提取sheet编码
            code_match = re.search(r"H7-\d+", name)
            if code_match:
                code = code_match.group()
                # 区分成本/公允版本
                if code in dual_mode_pairs:
                    if "成本" in name or "原值" in name:
                        key_name = f"{code}_cost"
                    elif "公允" in name:
                        key_name = f"{code}_fair"
                    else:
                        key_name = code
                else:
                    key_name = code
            else:
                key_name = name

            key_sheets[key_name] = {
                "title": name,
                "size": info["size"],
                "total_formula_count": info["formula_count"],
                "business_formula_count": info["business_formula_count"],
                "header_ref_count": info["header_ref_count"],
                "formula_patterns": info["formula_patterns"],
                "business_formula_rows": info["business_formula_rows"],
            }

    # 构建最终JSON
    result = {
        "source_file": "H7 生产性生物资产.xlsx",
        "source_path": str(xlsx_path.relative_to(PROJECT_ROOT)),
        "total_sheets": total_sheets,
        "sheet_names": sheet_names,
        "skip_sheets": skip_sheets,
        "active_sheets": active_sheets,
        "active_sheet_count": len(active_sheets),
        "dual_mode_pairs": dual_mode_pairs,
        "dual_mode_pair_codes": ["H7-1", "H7-2", "H7-6", "H7-7"],
        "key_sheets": key_sheets,
        "sheets": sheets_data,
        "validation": {
            "expected_total_sheets": 26,
            "actual_total_sheets": total_sheets,
            "sheet_count_match": total_sheets == 26,
            "dual_mode_identified": all(
                v["cost"] is not None and v["fair"] is not None
                for v in dual_mode_pairs.values()
            ),
            "dual_mode_details": dual_mode_pairs,
        },
    }

    # 写入JSON
    OUTPUT_PATH.parent.mkdir(parents=True, exist_ok=True)
    with open(OUTPUT_PATH, "w", encoding="utf-8") as f:
        json.dump(result, f, ensure_ascii=False, indent=2, cls=DateTimeEncoder)

    print(f"\n输出已写入: {OUTPUT_PATH}")
    print(f"\n=== 验证结果 ===")
    print(f"总sheet数: {total_sheets} (期望26)")
    print(f"活动sheet数: {len(active_sheets)}")
    print(f"跳过sheet数: {len(skip_sheets)}")
    print(f"含公式key_sheets: {len(key_sheets)}")
    print(f"双模式对识别: {result['validation']['dual_mode_identified']}")
    for code, pair in dual_mode_pairs.items():
        print(f"  {code}: 成本={pair['cost']} | 公允={pair['fair']}")

    if not result["validation"]["sheet_count_match"]:
        print(f"\n⚠️ 警告: sheet数量不匹配! 实际={total_sheets}, 期望=26")
        print("  实际sheet列表:")
        for i, name in enumerate(sheet_names, 1):
            print(f"    {i}. {name}")

    wb.close()
    return result


if __name__ == "__main__":
    main()

"""
Task 0.1: openpyxl脚本读取N2应交税费.xlsx全部18 sheet
提取：sheet名/列头/行数/公式单元格/合并区域/数据类型
标记skip：O1A原底稿/出口退税额复核示例
产出：n2_structure_summary.json（权威列头+公式清单）
"""

import json
import re
from pathlib import Path
from openpyxl import load_workbook
from openpyxl.utils import get_column_letter

# --- Configuration ---
SOURCE_FILE = Path(r"d:\GT_plan\数据\致同通用审计程序及底稿模板（2025年修订）\1.致同审计程序及底稿模板（2025年）\4.风险应对-实质性程序（D-N）\N 税金循环\N2 应交税费.xlsx")
OUTPUT_FILE = Path(r"d:\GT_plan\.kiro\specs\n2-taxes-payable\n2_structure_summary.json")

# Sheets to mark as skip (auxiliary/reference only)
SKIP_PATTERNS = ["O1A", "出口退税额复核示例"]

# Key sheets to confirm dimensions (from task description)
KEY_SHEET_EXPECTATIONS = {
    "N2-10": {"expected_size_hint": "51×7", "expected_formulas_hint": None},
    "N2-11": {"expected_size_hint": "26×16", "expected_formulas_hint": None},
    "N2-1": {"expected_size_hint": "31×14", "expected_formulas_hint": 85},
    "N2-2": {"expected_size_hint": "42×23", "expected_formulas_hint": 22},
    "N2-5": {"expected_size_hint": "54×15", "expected_formulas_hint": None},
    "N2-6": {"expected_size_hint": "48×8", "expected_formulas_hint": None},
    "N2-8": {"expected_size_hint": "26×9", "expected_formulas_hint": 11},
    "N2-9": {"expected_size_hint": "30×7", "expected_formulas_hint": None},
}


def is_skip_sheet(title: str) -> bool:
    """Check if a sheet should be marked as skip."""
    for pattern in SKIP_PATTERNS:
        if pattern in title:
            return True
    return False


def is_header_ref_formula(formula: str) -> bool:
    """Check if formula is a header reference (e.g., =底稿目录!A2)."""
    return bool(re.match(r"^=底稿目录!", formula))


def anonymize_formula(formula: str) -> str:
    """Replace cell references with REF for pattern deduplication."""
    # Replace sheet references like '递延所得税资产明细表N1-2'!$B$11
    anon = re.sub(r"'[^']+?'!", "'REF'!", formula)
    # Replace cell refs like $A$7, A7, $A7, A$7
    anon = re.sub(r"\$?[A-Z]{1,3}\$?\d+", "REF", anon)
    return anon


def extract_sheet_info(ws) -> dict:
    """Extract comprehensive info from a single worksheet."""
    title = ws.title
    dims = ws.dimensions
    max_row = ws.max_row
    max_col = ws.max_column
    size = f"{max_row}×{max_col}"

    # Merged cells
    merged = [str(m) for m in ws.merged_cells.ranges]

    # Headers (first 3 rows)
    headers = {}
    for row_idx in range(1, min(4, max_row + 1)):
        row_data = []
        for col_idx in range(1, max_col + 1):
            cell = ws.cell(row=row_idx, column=col_idx)
            val = cell.value
            if val is not None:
                col_letter = get_column_letter(col_idx)
                if isinstance(val, str) and val.startswith("="):
                    row_data.append({"col": col_letter, "value": val})
                else:
                    row_data.append({"col": col_letter, "value": str(val)})
        if row_data:
            headers[f"row{row_idx}"] = row_data

    # Formula cells and data types
    formula_cells = []
    business_formula_cells = []
    data_types = {}
    header_ref_count = 0

    for row in ws.iter_rows(min_row=1, max_row=max_row, max_col=max_col):
        for cell in row:
            # Data type stats
            dt = cell.data_type
            data_types[dt] = data_types.get(dt, 0) + 1

            # Formula extraction
            if cell.data_type == "f" or (isinstance(cell.value, str) and cell.value.startswith("=")):
                formula_str = cell.value if isinstance(cell.value, str) else str(cell.value)
                col_letter = get_column_letter(cell.column)
                entry = {"cell": f"{col_letter}{cell.row}", "formula": formula_str}
                formula_cells.append(entry)

                if is_header_ref_formula(formula_str):
                    header_ref_count += 1
                else:
                    business_formula_cells.append(entry)

    # Unique formula patterns (anonymized)
    pattern_set = set()
    for fc in business_formula_cells:
        pattern_set.add(anonymize_formula(fc["formula"]))
    formula_patterns = sorted(pattern_set)[:20]  # Cap at 20 patterns

    # Business formula rows (unique rows with business formulas)
    biz_formula_rows = len(set(
        int(re.search(r"\d+", fc["cell"]).group())
        for fc in business_formula_cells
    )) if business_formula_cells else 0

    # Sample data rows (rows 3-10 with actual data)
    sample_rows = []
    for row_idx in range(3, min(11, max_row + 1)):
        row_data = {}
        for col_idx in range(1, max_col + 1):
            cell = ws.cell(row=row_idx, column=col_idx)
            val = cell.value
            if val is not None:
                col_letter = get_column_letter(col_idx)
                if isinstance(val, str) and val.startswith("="):
                    row_data[col_letter] = f"[FORMULA] {val}"
                else:
                    row_data[col_letter] = str(val) if not isinstance(val, (int, float)) else val
        if row_data:
            sample_rows.append({"row": row_idx, "data": row_data})

    return {
        "title": title,
        "dimensions": dims,
        "max_row": max_row,
        "max_column": max_col,
        "size": size,
        "merged_cells": merged,
        "merged_count": len(merged),
        "headers": headers,
        "formula_cells": formula_cells,
        "formula_count": len(formula_cells),
        "business_formula_cells": business_formula_cells,
        "business_formula_count": len(business_formula_cells),
        "header_ref_count": header_ref_count,
        "formula_patterns": formula_patterns,
        "sample_data_rows": sample_rows,
        "data_types": data_types,
        "business_formula_rows": biz_formula_rows,
    }


def identify_key_sheet(title: str) -> str | None:
    """Try to identify a key sheet code from title."""
    # Sort by length descending to match N2-10 before N2-1
    for code in sorted(KEY_SHEET_EXPECTATIONS.keys(), key=len, reverse=True):
        if code in title:
            return code
    return None


def main():
    print(f"Reading: {SOURCE_FILE}")
    if not SOURCE_FILE.exists():
        # Try alternate paths
        alt_paths = [
            Path(r"d:\GT_plan\backend\wp_templates\N\N2 应交税费.xlsx"),
            Path(r"d:\GT_plan\backend\storage\knowledge\03511cc3-f4b3-4a0f-ab6b-10160c176fcc\N2 应交税费.xlsx"),
        ]
        for alt in alt_paths:
            if alt.exists():
                print(f"  Using alternate: {alt}")
                source = alt
                break
        else:
            print("ERROR: Cannot find N2 应交税费.xlsx in any known location!")
            return
    else:
        source = SOURCE_FILE

    wb = load_workbook(source, data_only=False, read_only=False)

    result = {
        "source_file": "N2 应交税费.xlsx",
        "total_sheets": len(wb.sheetnames),
        "sheet_names": wb.sheetnames,
        "skip_sheets": [],
        "active_sheets": [],
        "active_sheet_count": 0,
        "key_sheets": {},
        "sheets": [],
    }

    for title in wb.sheetnames:
        if is_skip_sheet(title):
            result["skip_sheets"].append(title)
        else:
            result["active_sheets"].append(title)

    result["active_sheet_count"] = len(result["active_sheets"])

    print(f"Total sheets: {result['total_sheets']}")
    print(f"Active sheets: {result['active_sheet_count']}")
    print(f"Skip sheets: {result['skip_sheets']}")
    print()

    for title in wb.sheetnames:
        ws = wb[title]
        skip = is_skip_sheet(title)

        if skip:
            # Still extract basic info for skip sheets
            info = {
                "title": title,
                "skip": True,
                "dimensions": ws.dimensions,
                "max_row": ws.max_row,
                "max_column": ws.max_column,
                "size": f"{ws.max_row}×{ws.max_column}",
                "reason": "辅助sheet/参考底稿，不做HTML组件化",
            }
            result["sheets"].append(info)
            print(f"  [SKIP] {title} ({info['size']})")
        else:
            info = extract_sheet_info(ws)
            result["sheets"].append(info)

            # Check if it's a key sheet
            code = identify_key_sheet(title)
            if code:
                key_entry = {
                    "title": title,
                    "size": info["size"],
                    "total_formula_count": info["formula_count"],
                    "business_formula_count": info["business_formula_count"],
                    "header_ref_count": info["header_ref_count"],
                    "formula_patterns": info["formula_patterns"],
                    "business_formula_rows": info["business_formula_rows"],
                }
                # Add expected vs actual comparison
                exp = KEY_SHEET_EXPECTATIONS[code]
                key_entry["expected_size"] = exp["expected_size_hint"]
                key_entry["expected_formulas"] = exp["expected_formulas_hint"]
                result["key_sheets"][f"{code}_{title.split('N2')[0].strip() if 'N2' not in title[:5] else title.replace('N2', '').strip()}"] = key_entry

                status = "✓" if True else "✗"
                print(f"  [KEY] {title}: {info['size']}, {info['business_formula_count']} biz formulas (expected ~{exp['expected_formulas_hint']})")
            else:
                print(f"  {title}: {info['size']}, {info['formula_count']} formulas")

    # Rebuild key_sheets with cleaner keys
    clean_keys = {}
    for title in wb.sheetnames:
        code = identify_key_sheet(title)
        if code and not is_skip_sheet(title):
            ws_info = next((s for s in result["sheets"] if s["title"] == title), None)
            if ws_info and "formula_count" in ws_info:
                suffix = title.replace(f"N2", "").strip(" -")
                clean_key = f"{code}_{suffix[:6]}" if suffix else code
                # Use simpler key format
                for k, v in result["key_sheets"].items():
                    if v["title"] == title:
                        clean_keys[code] = v
                        break
    if clean_keys:
        result["key_sheets"] = clean_keys

    # Write output
    OUTPUT_FILE.parent.mkdir(parents=True, exist_ok=True)
    with open(OUTPUT_FILE, "w", encoding="utf-8") as f:
        json.dump(result, f, ensure_ascii=False, indent=2)

    print(f"\nOutput written to: {OUTPUT_FILE}")
    print(f"Key sheets confirmed: {list(result['key_sheets'].keys())}")

    wb.close()


if __name__ == "__main__":
    main()

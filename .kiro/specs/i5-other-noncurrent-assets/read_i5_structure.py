"""
Phase0 Task 0.1: 读取 I5 其他非流动资产.xlsx 全部 sheet 结构
产出: i5_structure_summary.json
"""
import json
import re
from pathlib import Path
from openpyxl import load_workbook
from openpyxl.utils import get_column_letter

# 双源: 优先用原始模板源
SOURCE_PATH = Path(
    r"D:\GT_plan\基础数据\致同通用审计程序及底稿模板（2025年修订）"
    r"\1.致同审计程序及底稿模板（2025年）\4.风险应对-实质性程序（D-N）"
    r"\I 无形资产循环\I5 其他非流动资产.xlsx"
)
OUTPUT_PATH = Path(r"D:\GT_plan\.kiro\specs\i5-other-noncurrent-assets\i5_structure_summary.json")


def extract_sheet_info(ws):
    """提取单个 sheet 的结构信息"""
    info = {
        "sheet_name": ws.title,
        "dimensions": ws.dimensions,
        "max_row": ws.max_row,
        "max_column": ws.max_column,
        "column_count": ws.max_column,
        "row_count": ws.max_row,
    }

    # 合并单元格
    merged = []
    for mc in ws.merged_cells.ranges:
        merged.append(str(mc))
    info["merged_cells_count"] = len(merged)
    info["merged_cells"] = merged[:30]  # 只记录前30个

    # 前5行内容（表头区域）
    header_rows = []
    for row_idx in range(1, min(6, ws.max_row + 1)):
        row_data = []
        for col_idx in range(1, ws.max_column + 1):
            cell = ws.cell(row=row_idx, column=col_idx)
            val = cell.value
            if val is not None:
                row_data.append({
                    "col": get_column_letter(col_idx),
                    "value": str(val)[:100],
                    "is_formula": str(val).startswith("=") if isinstance(val, str) else False,
                })
        if row_data:
            header_rows.append({"row": row_idx, "cells": row_data})
    info["header_rows"] = header_rows

    # 公式统计
    formulas = []
    formula_patterns = {}
    for row in ws.iter_rows(min_row=1, max_row=ws.max_row, max_col=ws.max_column):
        for cell in row:
            if isinstance(cell.value, str) and cell.value.startswith("="):
                coord = f"{get_column_letter(cell.column)}{cell.row}"
                formula = cell.value
                formulas.append({"cell": coord, "formula": formula[:200]})
                # 提取公式模式（去掉行号）
                pattern = re.sub(r"\d+", "N", formula)
                formula_patterns[pattern] = formula_patterns.get(pattern, 0) + 1

    info["formula_count"] = len(formulas)
    info["formulas"] = formulas[:50]  # 前50个公式
    info["formula_patterns"] = dict(
        sorted(formula_patterns.items(), key=lambda x: -x[1])[:20]
    )

    # 列宽/隐藏列
    hidden_cols = []
    col_widths = {}
    for col_idx in range(1, ws.max_column + 1):
        letter = get_column_letter(col_idx)
        dim = ws.column_dimensions.get(letter)
        if dim:
            if dim.hidden:
                hidden_cols.append(letter)
            if dim.width:
                col_widths[letter] = round(dim.width, 1)
    info["hidden_columns"] = hidden_cols
    info["column_widths"] = col_widths

    # 隐藏行
    hidden_rows = []
    for row_idx in range(1, ws.max_row + 1):
        dim = ws.row_dimensions.get(row_idx)
        if dim and dim.hidden:
            hidden_rows.append(row_idx)
    info["hidden_rows"] = hidden_rows

    return info


def main():
    print(f"Reading: {SOURCE_PATH}")
    assert SOURCE_PATH.exists(), f"File not found: {SOURCE_PATH}"

    wb = load_workbook(str(SOURCE_PATH), data_only=False)
    print(f"Sheet count: {len(wb.sheetnames)}")
    print(f"Sheet names: {wb.sheetnames}")

    result = {
        "source_file": str(SOURCE_PATH),
        "sheet_count": len(wb.sheetnames),
        "sheet_names": wb.sheetnames,
        "科目": "1911 其他非流动资产",
        "科目方向": "借方/资产类",
        "sheets": []
    }

    for ws_name in wb.sheetnames:
        ws = wb[ws_name]
        print(f"  Processing: {ws_name} ({ws.max_row}行 x {ws.max_column}列)")
        sheet_info = extract_sheet_info(ws)
        result["sheets"].append(sheet_info)

    # 汇总
    total_formulas = sum(s["formula_count"] for s in result["sheets"])
    result["total_formula_count"] = total_formulas
    result["summary"] = {
        "total_sheets": len(result["sheets"]),
        "total_formulas": total_formulas,
        "sheets_overview": [
            {
                "name": s["sheet_name"],
                "rows": s["max_row"],
                "cols": s["max_column"],
                "formulas": s["formula_count"],
                "merged_cells": s["merged_cells_count"]
            }
            for s in result["sheets"]
        ]
    }

    OUTPUT_PATH.parent.mkdir(parents=True, exist_ok=True)
    with open(OUTPUT_PATH, "w", encoding="utf-8") as f:
        json.dump(result, f, ensure_ascii=False, indent=2)

    print(f"\nOutput: {OUTPUT_PATH}")
    print(f"Total formulas: {total_formulas}")
    print("\nSheet overview:")
    for s in result["summary"]["sheets_overview"]:
        print(f"  {s['name']:30s}  {s['rows']:3d}行 {s['cols']:2d}列 {s['formulas']:3d}公式 {s['merged_cells']}合并")


if __name__ == "__main__":
    main()

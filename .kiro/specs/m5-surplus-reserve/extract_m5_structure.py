"""
M5盈余公积.xlsx 结构提取脚本
Phase 0 Task 0.1: openpyxl读取全部10 sheet，提取结构+确认公式数
产出：m5_structure_summary.json
"""
import json
import os
from openpyxl import load_workbook
from openpyxl.utils import get_column_letter

# 源模板路径
XLSX_PATH = r"d:\GT_plan\数据\致同通用审计程序及底稿模板（2025年修订）\1.致同审计程序及底稿模板（2025年）\4.风险应对-实质性程序（D-N）\M 权益循环\M5 盈余公积.xlsx"
OUTPUT_PATH = r"d:\GT_plan\.kiro\specs\m5-surplus-reserve\m5_structure_summary.json"


def extract_formulas(ws):
    """提取sheet中所有公式单元格"""
    formulas = []
    for row in ws.iter_rows():
        for cell in row:
            if cell.value and isinstance(cell.value, str) and cell.value.startswith('='):
                formulas.append({
                    "cell": f"{get_column_letter(cell.column)}{cell.row}",
                    "formula": cell.value
                })
    return formulas


def extract_headers(ws, max_header_rows=3):
    """提取前几行作为表头"""
    headers = []
    for row_idx in range(1, min(max_header_rows + 1, ws.max_row + 1)):
        row_data = []
        for col_idx in range(1, ws.max_column + 1):
            val = ws.cell(row=row_idx, column=col_idx).value
            if val is not None:
                row_data.append({
                    "col": get_column_letter(col_idx),
                    "value": str(val)[:100]  # 截断长文本
                })
        if row_data:
            headers.append({"row": row_idx, "cells": row_data})
    return headers


def extract_merged_cells(ws):
    """提取合并单元格信息"""
    merged = []
    for mc in ws.merged_cells.ranges:
        merged.append(str(mc))
    return merged


def analyze_sheet(ws, sheet_name):
    """分析单个sheet的完整结构"""
    formulas = extract_formulas(ws)
    headers = extract_headers(ws, max_header_rows=5)
    merged = extract_merged_cells(ws)

    # 非空行列统计
    non_empty_rows = set()
    non_empty_cols = set()
    for row in ws.iter_rows():
        for cell in row:
            if cell.value is not None:
                non_empty_rows.add(cell.row)
                non_empty_cols.add(cell.column)

    return {
        "sheet_name": sheet_name,
        "dimensions": ws.dimensions,
        "max_row": ws.max_row,
        "max_column": ws.max_column,
        "non_empty_rows": len(non_empty_rows),
        "non_empty_cols": len(non_empty_cols),
        "formula_count": len(formulas),
        "formulas": formulas,
        "headers": headers,
        "merged_cells_count": len(merged),
        "merged_cells": merged[:30],  # 只保留前30个
    }


def main():
    print(f"正在读取: {XLSX_PATH}")
    if not os.path.exists(XLSX_PATH):
        print(f"ERROR: 文件不存在 {XLSX_PATH}")
        return

    wb = load_workbook(XLSX_PATH, data_only=False)  # data_only=False保留公式
    print(f"工作簿sheet数: {len(wb.sheetnames)}")
    print(f"Sheet名称: {wb.sheetnames}")

    summary = {
        "source_file": "M5 盈余公积.xlsx",
        "source_path": XLSX_PATH,
        "total_sheets": len(wb.sheetnames),
        "sheet_names": wb.sheetnames,
        "sheets": {},
        "formula_summary": {},
        "validation": {}
    }

    total_formulas = 0
    for sheet_name in wb.sheetnames:
        ws = wb[sheet_name]
        print(f"\n--- 分析 sheet: {sheet_name} ---")
        sheet_info = analyze_sheet(ws, sheet_name)
        summary["sheets"][sheet_name] = sheet_info
        total_formulas += sheet_info["formula_count"]
        print(f"  维度: {sheet_info['dimensions']}, 行×列: {sheet_info['max_row']}×{sheet_info['max_column']}")
        print(f"  非空行: {sheet_info['non_empty_rows']}, 非空列: {sheet_info['non_empty_cols']}")
        print(f"  公式数: {sheet_info['formula_count']}")
        print(f"  合并区域数: {sheet_info['merged_cells_count']}")

    summary["formula_summary"] = {
        "total_formulas": total_formulas,
        "by_sheet": {name: summary["sheets"][name]["formula_count"] for name in wb.sheetnames}
    }

    # 验证：确认审定表M5-1(57公式) + 明细M5-2(12公式) + 计提检查M5-4(11公式)
    validation = {}
    for sheet_name, info in summary["sheets"].items():
        fc = info["formula_count"]
        if "M5-1" in sheet_name or "审定" in sheet_name:
            validation["M5-1_审定表"] = {
                "sheet_name": sheet_name,
                "formula_count": fc,
                "expected": 57,
                "match": fc == 57
            }
        elif "M5-2" in sheet_name or "明细" in sheet_name:
            validation["M5-2_明细表"] = {
                "sheet_name": sheet_name,
                "formula_count": fc,
                "expected": 12,
                "match": fc == 12
            }
        elif "M5-4" in sheet_name or "计提" in sheet_name:
            validation["M5-4_计提检查"] = {
                "sheet_name": sheet_name,
                "formula_count": fc,
                "expected": 11,
                "match": fc == 11
            }

    summary["validation"] = validation

    # 写出JSON
    with open(OUTPUT_PATH, 'w', encoding='utf-8') as f:
        json.dump(summary, f, ensure_ascii=False, indent=2)

    print(f"\n{'='*60}")
    print(f"总计公式: {total_formulas}")
    print(f"验证结果: {json.dumps(validation, ensure_ascii=False, indent=2)}")
    print(f"\n输出已写入: {OUTPUT_PATH}")


if __name__ == "__main__":
    main()

"""
Phase 0 Task 0.1: openpyxl脚本读取L8财务费用.xlsx全部10 sheet
提取结构 + 确认明细表L8-2(37公式)+L8-4非金融利息+L8-5截止测试结构
产出：l8_structure_summary.json
"""
import json
import os
from openpyxl import load_workbook
from openpyxl.utils import get_column_letter

# Source xlsx path
XLSX_PATH = r"d:\GT_plan\数据\致同通用审计程序及底稿模板（2025年修订）\1.致同审计程序及底稿模板（2025年）\4.风险应对-实质性程序（D-N）\L 债务循环\L8 财务费用.xlsx"
OUTPUT_PATH = r"d:\GT_plan\.kiro\specs\l8-financial-expenses\l8_structure_summary.json"


def extract_sheet_structure(ws):
    """Extract structure from a worksheet."""
    sheet_info = {
        "sheet_name": ws.title,
        "max_row": ws.max_row,
        "max_column": ws.max_column,
        "dimensions": ws.dimensions,
        "column_headers": [],
        "formulas": [],
        "formula_count": 0,
        "merged_cells": [str(m) for m in ws.merged_cells.ranges],
        "sample_data_rows": []
    }

    # Extract column headers (first few rows typically contain headers)
    for row_idx in range(1, min(6, ws.max_row + 1)):
        row_headers = []
        for col_idx in range(1, ws.max_column + 1):
            cell = ws.cell(row=row_idx, column=col_idx)
            val = cell.value
            if val is not None:
                row_headers.append({
                    "col": get_column_letter(col_idx),
                    "col_idx": col_idx,
                    "value": str(val)[:100]  # truncate long values
                })
        if row_headers:
            sheet_info["column_headers"].append({
                "row": row_idx,
                "cells": row_headers
            })

    # Scan all cells for formulas
    formula_cells = []
    for row in ws.iter_rows(min_row=1, max_row=ws.max_row, max_col=ws.max_column):
        for cell in row:
            if cell.value and isinstance(cell.value, str) and cell.value.startswith("="):
                formula_cells.append({
                    "cell": f"{get_column_letter(cell.column)}{cell.row}",
                    "formula": cell.value[:200]  # truncate very long formulas
                })

    sheet_info["formulas"] = formula_cells
    sheet_info["formula_count"] = len(formula_cells)

    # Sample data rows (rows 6-15 for context)
    for row_idx in range(6, min(16, ws.max_row + 1)):
        row_data = []
        for col_idx in range(1, min(ws.max_column + 1, 25)):  # limit to 24 cols
            cell = ws.cell(row=row_idx, column=col_idx)
            val = cell.value
            if val is not None:
                row_data.append({
                    "col": get_column_letter(col_idx),
                    "value": str(val)[:80]
                })
        if row_data:
            sheet_info["sample_data_rows"].append({
                "row": row_idx,
                "cells": row_data
            })

    return sheet_info


def main():
    print(f"Loading workbook: {XLSX_PATH}")
    if not os.path.exists(XLSX_PATH):
        print(f"ERROR: File not found: {XLSX_PATH}")
        return

    wb = load_workbook(XLSX_PATH, data_only=False)  # data_only=False to see formulas

    result = {
        "source_file": XLSX_PATH,
        "total_sheets": len(wb.sheetnames),
        "sheet_names": wb.sheetnames,
        "sheets": [],
        "key_confirmations": {}
    }

    print(f"Total sheets: {len(wb.sheetnames)}")
    print(f"Sheet names: {wb.sheetnames}")
    print()

    for ws_name in wb.sheetnames:
        ws = wb[ws_name]
        print(f"Processing: {ws_name} ({ws.max_row}×{ws.max_column})")
        info = extract_sheet_structure(ws)
        result["sheets"].append(info)

    # Key confirmations
    # L8-2 明细表 (37 formulas)
    l8_2_sheets = [s for s in result["sheets"] if "L8-2" in s["sheet_name"] or "明细" in s["sheet_name"]]
    if l8_2_sheets:
        l8_2 = l8_2_sheets[0]
        result["key_confirmations"]["L8-2_detail"] = {
            "sheet_name": l8_2["sheet_name"],
            "dimensions": f"{l8_2['max_row']}×{l8_2['max_column']}",
            "formula_count": l8_2["formula_count"],
            "expected_formulas": 37,
            "match": l8_2["formula_count"] == 37,
            "note": "明细表L8-2 公式数量确认"
        }

    # L8-4 非金融机构利息
    l8_4_sheets = [s for s in result["sheets"] if "L8-4" in s["sheet_name"] or "非金融" in s["sheet_name"]]
    if l8_4_sheets:
        l8_4 = l8_4_sheets[0]
        result["key_confirmations"]["L8-4_non_financial_interest"] = {
            "sheet_name": l8_4["sheet_name"],
            "dimensions": f"{l8_4['max_row']}×{l8_4['max_column']}",
            "formula_count": l8_4["formula_count"],
            "note": "非金融机构利息支出测算表结构确认"
        }

    # L8-5 截止测试
    l8_5_sheets = [s for s in result["sheets"] if "L8-5" in s["sheet_name"] or "截止" in s["sheet_name"]]
    if l8_5_sheets:
        l8_5 = l8_5_sheets[0]
        result["key_confirmations"]["L8-5_cutoff_test"] = {
            "sheet_name": l8_5["sheet_name"],
            "dimensions": f"{l8_5['max_row']}×{l8_5['max_column']}",
            "formula_count": l8_5["formula_count"],
            "note": "截止性测试结构确认"
        }

    # Summary statistics
    total_formulas = sum(s["formula_count"] for s in result["sheets"])
    result["summary"] = {
        "total_formulas_all_sheets": total_formulas,
        "sheets_with_formulas": [
            {"name": s["sheet_name"], "count": s["formula_count"]}
            for s in result["sheets"] if s["formula_count"] > 0
        ]
    }

    # Write output
    with open(OUTPUT_PATH, "w", encoding="utf-8") as f:
        json.dump(result, f, ensure_ascii=False, indent=2)

    print(f"\nOutput written to: {OUTPUT_PATH}")
    print(f"\nSummary:")
    print(f"  Total sheets: {result['total_sheets']}")
    print(f"  Total formulas: {total_formulas}")
    print(f"  Key confirmations:")
    for key, val in result["key_confirmations"].items():
        print(f"    {key}: {val['sheet_name']} ({val['dimensions']}, {val['formula_count']} formulas)")


if __name__ == "__main__":
    main()

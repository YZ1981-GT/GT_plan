"""
analyze_h1_structure.py
=======================
读取 H1 固定资产.xlsx 全部 sheet，提取结构信息并输出 JSON。

Phase 0 双源输入验证 - 任务 0.1

提取：sheet名 / 列头 / 行数 / 公式单元格 / 合并区域 / 数据类型
产出：.kiro/specs/h1-fixed-assets/h1_structure_summary.json
验证：26 sheet 结构与 spec requirements.md Glossary 描述一致

用法：python backend/scripts/analyze_h1_structure.py
"""

import json
import os
import sys
from pathlib import Path
from collections import Counter

import openpyxl
from openpyxl.utils import get_column_letter


# ---------- 配置 ----------

_CANDIDATES = [
    Path(
        r"基础数据/致同通用审计程序及底稿模板（2025年修订）"
        r"/1.致同审计程序及底稿模板（2025年）"
        r"/4.风险应对-实质性程序（D-N）/H 固定资产循环/H1 固定资产.xlsx"
    ),
]

OUTPUT_FILE = Path(".kiro/specs/h1-fixed-assets/h1_structure_summary.json")

# Glossary 中描述的 26 个 sheet 预期结构（名称关键词→预期特征）
EXPECTED_SHEETS = {
    "Tab_Index": {"keyword": "目录", "expected_cols_approx": 8},
    "Procedure_Table_H1A": {"keyword": "H1A", "expected_cols_approx": 13},
    "Adjudication_H1_1": {"keyword": "H1-1", "expected_cols_approx": 9},
    "Disclosure_Listed": {"keyword": "上市", "expected_cols_approx": 257},
    "Disclosure_SOE": {"keyword": "国有", "expected_cols_approx": 256},
    "Detail_H1_2": {"keyword": "H1-2", "expected_cols_approx": 54},
    "Adjustment_H1_3": {"keyword": "H1-3", "expected_cols_approx": 13},
    "Idle_Check_H1_4": {"keyword": "H1-4", "expected_cols_approx": 12},
    "Policy_Check_H1_5": {"keyword": "H1-5", "expected_cols_approx": 16},
    "Analysis_H1_6": {"keyword": "H1-6", "expected_cols_approx": 20},
    "Addition_Check_H1_7": {"keyword": "H1-7", "expected_cols_approx": 24},
    "Disposal_Check_H1_8": {"keyword": "H1-8", "expected_cols_approx": 27},
    "Stocktake_Plan_H1_9": {"keyword": "H1-9", "expected_cols_approx": 15},
    "Stocktake_Check_H1_10": {"keyword": "H1-10", "expected_cols_approx": 17},
    "Stocktake_Summary_H1_11": {"keyword": "H1-11", "expected_cols_approx": 10},
    "Depreciation_Straight_H1_12": {"keyword": "不含减值", "expected_cols_approx": 28},
    "Depreciation_Impair_H1_12": {"keyword": "折旧测算表（含减值）", "expected_cols_approx": 28},
    "Depreciation_Multi_H1_12": {"keyword": "多次减值", "expected_cols_approx": 29},
    "Depreciation_Alloc_H1_13": {"keyword": "H1-13", "expected_cols_approx": 11},
    "Impairment_H1_14": {"keyword": "H1-14", "expected_cols_approx": 33},
    "Recoverable_H1_15": {"keyword": "H1-15", "expected_cols_approx": 28},
    "Title_Building_H1_16": {"keyword": "H1-16", "expected_cols_approx": 22},
    "Title_Vehicle_H1_17": {"keyword": "H1-17", "expected_cols_approx": 18},
    "Related_Party_H1_18": {"keyword": "H1-18", "expected_cols_approx": 15},
    "Operating_Lease_H1_19": {"keyword": "H1-19", "expected_cols_approx": 25},
    "Finance_Lease_H1_20": {"keyword": "H1-20", "expected_cols_approx": 22},
}


def find_source_file() -> Path:
    """按候选列表查找第一个存在的 xlsx 文件。"""
    for candidate in _CANDIDATES:
        if candidate.exists():
            return candidate
    # 尝试绝对路径（从项目根）
    project_root = Path(os.getcwd())
    for candidate in _CANDIDATES:
        abs_path = project_root / candidate
        if abs_path.exists():
            return abs_path
    print("ERROR: 找不到 H1 固定资产.xlsx，尝试过：")
    for c in _CANDIDATES:
        print(f"  - {c}")
    sys.exit(1)


def extract_headers(ws, max_scan_rows: int = 5) -> dict:
    """从前几行提取列头信息，返回所有可能的列头行。"""
    header_rows = {}
    for row_idx in range(1, min(max_scan_rows + 1, ws.max_row + 1)):
        headers = []
        for col_idx in range(1, ws.max_column + 1):
            cell = ws.cell(row=row_idx, column=col_idx)
            val = cell.value
            if val is not None:
                headers.append(str(val).strip())
            else:
                headers.append("")
        non_empty = [h for h in headers if h]
        if len(non_empty) >= 2:
            header_rows[f"row_{row_idx}"] = non_empty
    return header_rows


def extract_primary_headers(ws, max_scan_rows: int = 5) -> list[str]:
    """提取最可能是主列头的一行（非空数最多的行）。"""
    best_row = []
    best_count = 0
    for row_idx in range(1, min(max_scan_rows + 1, ws.max_row + 1)):
        headers = []
        for col_idx in range(1, ws.max_column + 1):
            cell = ws.cell(row=row_idx, column=col_idx)
            val = cell.value
            if val is not None:
                headers.append(str(val).strip())
            else:
                headers.append("")
        non_empty = sum(1 for h in headers if h)
        if non_empty > best_count:
            best_count = non_empty
            best_row = [h for h in headers if h]
    return best_row


def extract_formulas(ws) -> dict:
    """提取公式单元格信息。"""
    formulas = []
    formula_count = 0
    formula_by_col = Counter()
    for row in ws.iter_rows(min_row=1, max_row=ws.max_row, max_col=ws.max_column):
        for cell in row:
            if cell.value and isinstance(cell.value, str) and cell.value.startswith("="):
                formula_count += 1
                col_letter = get_column_letter(cell.column)
                formula_by_col[col_letter] += 1
                if len(formulas) < 15:  # 保留前 15 个样例
                    formulas.append({
                        "cell": cell.coordinate,
                        "formula": cell.value,
                    })
    return {
        "count": formula_count,
        "samples": formulas,
        "by_column": dict(formula_by_col.most_common(10)),
    }


def extract_merged_areas(ws) -> list[str]:
    """提取合并单元格区域。"""
    return [str(r) for r in ws.merged_cells.ranges]


def extract_data_types(ws) -> dict:
    """统计各列数据类型分布（采样前 30 行数据区）。"""
    type_summary = {}
    data_start_row = min(4, ws.max_row)  # 跳过列头
    sample_end_row = min(data_start_row + 30, ws.max_row)

    for col_idx in range(1, ws.max_column + 1):
        col_letter = get_column_letter(col_idx)
        types = Counter()
        for row_idx in range(data_start_row, sample_end_row + 1):
            cell = ws.cell(row=row_idx, column=col_idx)
            val = cell.value
            if val is None:
                types["null"] += 1
            elif isinstance(val, str):
                if val.startswith("="):
                    types["formula"] += 1
                else:
                    types["string"] += 1
            elif isinstance(val, (int, float)):
                types["number"] += 1
            elif hasattr(val, "strftime"):
                types["date"] += 1
            else:
                types["other"] += 1

        # 只保留非空类型
        non_null_types = {k: v for k, v in types.items() if k != "null" and v > 0}
        if non_null_types:
            dominant_type = max(non_null_types, key=non_null_types.get)
            type_summary[col_letter] = dominant_type

    return type_summary


def validate_against_glossary(sheets: list[dict]) -> dict:
    """将实际 sheet 结构与 Glossary 预期进行对照验证。"""
    validation = {
        "total_sheets": len(sheets),
        "expected_sheets": 26,
        "match": len(sheets) == 26,
        "matched_glossary_entries": [],
        "unmatched_glossary_entries": [],
        "sheet_mapping": {},
    }

    sheet_names = [s["name"] for s in sheets]

    for glossary_key, expected in EXPECTED_SHEETS.items():
        keyword = expected["keyword"]
        matched_sheet = None
        for sname in sheet_names:
            if keyword in sname:
                matched_sheet = sname
                break

        if matched_sheet:
            validation["matched_glossary_entries"].append(glossary_key)
            validation["sheet_mapping"][glossary_key] = matched_sheet
        else:
            validation["unmatched_glossary_entries"].append(glossary_key)

    return validation


def main():
    # 切换到项目根目录
    project_root = Path(__file__).resolve().parent.parent.parent
    os.chdir(project_root)

    source_file = find_source_file()
    print(f"📂 读取文件: {source_file}")
    print(f"   文件大小: {source_file.stat().st_size / 1024:.1f} KB")

    # data_only=False 以读取公式而非计算结果
    wb = openpyxl.load_workbook(str(source_file), data_only=False, read_only=False)

    result = {
        "file_path": str(source_file),
        "file_size_kb": round(source_file.stat().st_size / 1024, 1),
        "sheet_count": len(wb.sheetnames),
        "sheet_names": wb.sheetnames,
        "sheets": [],
        "validation": None,
        "summary": {},
    }

    print(f"\n共 {len(wb.sheetnames)} 个 sheet:")
    print("-" * 70)

    total_formulas = 0
    total_merged = 0

    for idx, name in enumerate(wb.sheetnames, 1):
        ws = wb[name]
        print(f"  [{idx:2d}] {name}")
        print(f"       尺寸: {ws.max_row}行 × {ws.max_column}列")

        # 提取各项结构信息
        all_headers = extract_headers(ws)
        primary_headers = extract_primary_headers(ws)
        formula_info = extract_formulas(ws)
        merged_areas = extract_merged_areas(ws)
        data_types = extract_data_types(ws)

        total_formulas += formula_info["count"]
        total_merged += len(merged_areas)

        print(f"       列头: {len(primary_headers)}列 | 公式: {formula_info['count']}个 | 合并: {len(merged_areas)}个")

        sheet_info = {
            "index": idx,
            "name": name,
            "rows": ws.max_row,
            "cols": ws.max_column,
            "dimensions": f"{ws.max_row}行 × {ws.max_column}列",
            "primary_headers": primary_headers,
            "all_header_rows": all_headers,
            "formula_count": formula_info["count"],
            "formula_samples": formula_info["samples"],
            "formula_by_column": formula_info["by_column"],
            "merged_areas": merged_areas,
            "merged_area_count": len(merged_areas),
            "data_types": data_types,
        }
        result["sheets"].append(sheet_info)

    # 验证
    print("\n" + "=" * 70)
    print("📋 Glossary 对照验证:")
    validation = validate_against_glossary(result["sheets"])
    result["validation"] = validation

    if validation["match"]:
        print(f"  ✅ sheet 数量匹配: {validation['total_sheets']}/{validation['expected_sheets']}")
    else:
        print(f"  ⚠️ sheet 数量不匹配: 实际 {validation['total_sheets']} vs 预期 {validation['expected_sheets']}")

    print(f"  已匹配 Glossary 条目: {len(validation['matched_glossary_entries'])}/{len(EXPECTED_SHEETS)}")
    if validation["unmatched_glossary_entries"]:
        print(f"  ❌ 未匹配条目: {validation['unmatched_glossary_entries']}")

    # 汇总
    result["summary"] = {
        "total_sheets": len(wb.sheetnames),
        "total_formulas": total_formulas,
        "total_merged_areas": total_merged,
        "total_rows": sum(s["rows"] for s in result["sheets"]),
        "total_cols_max": max(s["cols"] for s in result["sheets"]),
        "avg_rows_per_sheet": round(sum(s["rows"] for s in result["sheets"]) / len(result["sheets"]), 1),
    }

    # 输出
    OUTPUT_FILE.parent.mkdir(parents=True, exist_ok=True)
    with open(OUTPUT_FILE, "w", encoding="utf-8") as f:
        json.dump(result, f, ensure_ascii=False, indent=2)

    print(f"\n✅ 产出文件: {OUTPUT_FILE}")
    print(f"   sheet 数量: {result['sheet_count']}")
    print(f"   公式总数: {total_formulas}")
    print(f"   合并区域总数: {total_merged}")
    print(f"   数据行总数: {result['summary']['total_rows']}")

    # 验证 26 sheet
    if result["sheet_count"] != 26:
        print(f"\n⚠️  注意：实际 sheet 数量为 {result['sheet_count']}，spec 描述为 26。")
        print("   可能存在隐藏 sheet 或 sheet 数量描述需要修正。")

    wb.close()
    return result


if __name__ == "__main__":
    main()

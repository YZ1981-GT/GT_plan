"""Repair mismatched F2 OnlyOffice caches (wrong pack / whole Excel)."""
from __future__ import annotations

import shutil
from pathlib import Path

from openpyxl import load_workbook

from app.routers.wp_onlyoffice_router import (
    _extract_sheet_code,
    _hide_non_target_sheets,
    _oo_cache_mismatches_template,
    _resolve_whole_workbook_file,
)
from app.services import wp_template_finder as f
from app.services.wp_template_finder import find_template_file_any, find_whole_workbook_template

f._index_cache = None

LABELS = {
    "F2": "存货实质性程序表F2A",
    "F2A": "存货实质性程序表F2A",
    "F2-1": "存货审定表F2-1",
    "F2-2": "明细汇总表F2-2",
    "F2-14": "调整分录汇总F2-14",
    "F2-16": "会计政策、核算流程F2-16",
}

project_ids: list[str] = []
root = Path("storage/projects")
if root.exists():
    project_ids.extend(p.name for p in root.iterdir() if p.is_dir())

for pid in project_ids:
    base = Path("storage/projects") / pid / "workpapers" / "onlyoffice"
    if not base.exists():
        continue
    print("=== project", pid, "===")

    whole_tpl = find_whole_workbook_template("F2")
    if whole_tpl:
        try:
            from uuid import UUID

            path = _resolve_whole_workbook_file(UUID(pid), "F2")
            print("whole", path.name, "<-", whole_tpl.name, path.stat().st_size)
        except Exception as e:
            print("whole FAIL", e)

    for p in sorted(base.glob("F2*.xlsx")):
        if p.name.endswith("__whole.xlsx"):
            continue
        stem = p.stem
        tpl = find_template_file_any(stem)
        if not tpl:
            print("skip no tpl", stem)
            continue
        label = LABELS.get(stem)
        need = _oo_cache_mismatches_template(p, wp_code=stem, visible_sheet=label)
        if not need:
            names = load_workbook(p, read_only=True).sheetnames
            codes = {_extract_sheet_code(n) for n in names}
            if stem not in ("F2",) and stem not in codes and not any(
                (c or "").startswith(stem) for c in codes if c
            ):
                # F2A / F2-2 等段码必须出现在 workbook 中
                need = True
            elif label and not any(label in n or n.endswith(stem) for n in names):
                need = True
        if need:
            shutil.copy2(tpl, p)
            if label:
                _hide_non_target_sheets(p, label)
            print("fixed", stem, "<-", tpl.name, p.stat().st_size)
        else:
            if label:
                _hide_non_target_sheets(p, label)
            print("keep", stem, p.stat().st_size)

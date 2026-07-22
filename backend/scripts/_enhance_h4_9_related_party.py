# -*- coding: utf-8 -*-
"""Enhance H4-9 关联交易检查表 Excel template."""
from __future__ import annotations

from pathlib import Path

import openpyxl
from openpyxl.comments import Comment
from openpyxl.styles import Alignment, Font, PatternFill, Border, Side
from openpyxl.worksheet.datavalidation import DataValidation

ROOT = Path(r"D:/GT_plan")
TARGETS = [ROOT / "backend" / "wp_templates" / "H" / "H4 工程物资.xlsx"]
base = ROOT / "基础数据"
if base.exists():
    for p in base.rglob("H4 工程物资.xlsx"):
        TARGETS.append(p)

SHEET = "关联交易检查表H4-9"
THIN = Border(
    left=Side(style="thin", color="B0B0B0"),
    right=Side(style="thin", color="B0B0B0"),
    top=Side(style="thin", color="B0B0B0"),
    bottom=Side(style="thin", color="B0B0B0"),
)
TOTAL_FILL = PatternFill("solid", fgColor="FFF2CC")
GUIDE_FONT = Font(color="0563C1", size=9)
BOLD = Font(bold=True, size=10)


def enhance(ws):
    if not str(ws["A8"].value or "").startswith("二、审计过程"):
        ws["A8"] = "二、审计过程："

    tip = (
        "编制提示：仅登记合并范围外关联方；购入关注入账价值与价款差异；"
        "出售净值=原值−减值；占同类%可按本表同类合计估算；"
        "价差超10%或标记异常须说明公允依据；无交易时在审计说明注明「本期无此类交易」。"
    )
    # 放在关联关系列表下方空行，避免占用数据验证源区
    tip_row = None
    for r in range(40, 60):
        if ws.cell(r, 1).value is None and tip_row is None:
            # 找到列表结束后的首个空行
            prev = ws.cell(r - 1, 1).value
            if prev in ("其他关联方", "关联关系") or (isinstance(prev, str) and "关联" in prev):
                tip_row = r
                break
    if tip_row is None:
        tip_row = 44
    if ws.cell(tip_row, 1).value is None:
        ws.cell(tip_row, 1).value = tip
        ws.cell(tip_row, 1).font = GUIDE_FONT
        ws.cell(tip_row, 1).alignment = Alignment(wrap_text=True, vertical="top")
        ws.row_dimensions[tip_row].height = 36

    for r in range(11, 17):
        for c in range(1, 13):
            cell = ws.cell(r, c)
            if cell.border.left.style is None:
                cell.border = THIN

    # purchase subtotal — only if empty
    if ws["A17"].value is None:
        ws["A17"] = "合计"
        ws["A17"].font = BOLD
        ws["E17"] = "=SUM(E11:E16)"
        ws["F17"] = "=SUM(F11:F16)"
        for c in range(1, 13):
            cell = ws.cell(17, c)
            cell.fill = TOTAL_FILL
            cell.border = THIN
            if c in (5, 6):
                cell.number_format = "#,##0.00"

    ws["E10"].comment = Comment("购买价款（与定价政策一致口径）", "H4-9")
    ws["F10"].comment = Comment(
        "工程物资入账价值。与购买价款差异通常为运杂/税费/折扣等，重大差异须在备注说明。",
        "H4-9",
    )
    ws["G10"].comment = Comment("占同类交易金额的比例% = 本笔价款 / 同类交易总额 × 100", "H4-9")
    ws["J10"].comment = Comment("填「是/否/待定」。选「是」时须在备注说明异常性质及追加程序。", "H4-9")

    for r in range(20, 24):
        ws.cell(r, 7).value = f'=IF(OR(E{r}<>"",F{r}<>""),E{r}-F{r},"")'
        ws.cell(r, 7).number_format = "#,##0.00"
        for c in range(1, 15):
            cell = ws.cell(r, c)
            if cell.border.left.style is None:
                cell.border = THIN

    ws["E19"].comment = Comment("出售时工程物资账面原值", "H4-9")
    ws["F19"].comment = Comment("出售时已计提减值准备", "H4-9")
    ws["G19"].comment = Comment("公式：出售时净值 = 原值 − 减值准备", "H4-9")
    ws["H19"].comment = Comment("销售价格（不含税）。可与净值比较评估处置损益。", "H4-9")
    ws["L19"].comment = Comment("填「是/否/待定」。选「是」时须在备注说明。", "H4-9")

    # insert sale total before 三、审计说明 if not already done
    if str(ws["A24"].value or "").startswith("三、"):
        ws.insert_rows(24)
        ws["A24"] = "合计"
        ws["A24"].font = BOLD
        ws["E24"] = "=SUM(E20:E23)"
        ws["F24"] = "=SUM(F20:F23)"
        ws["G24"] = "=SUM(G20:G23)"
        ws["H24"] = "=SUM(H20:H23)"
        for c in range(1, 15):
            cell = ws.cell(24, c)
            cell.fill = TOTAL_FILL
            cell.border = THIN
            if c in (5, 6, 7, 8):
                cell.number_format = "#,##0.00"

    # avoid duplicate validations on re-run
    existing = []
    for dv in list(ws.data_validations.dataValidation):
        existing.append(str(dv.sqref))
    if "J11:J16" not in "".join(existing):
        dv = DataValidation(type="list", formula1='"是,否,待定"', allow_blank=True)
        dv.error = "请选择：是 / 否 / 待定"
        dv.errorTitle = "关联交易是否存在异常"
        ws.add_data_validation(dv)
        dv.add("J11:J16")
        dv.add("L20:L23")

    for r in range(30, 55):
        if ws.cell(r, 1).value == "关联关系":
            note_r = r - 1
            if ws.cell(note_r, 2).value is None:
                ws.cell(note_r, 2).value = "（下拉选项；本表仅适用于合并范围外关联方）"
                ws.cell(note_r, 2).font = GUIDE_FONT
            break

    for r in range(25, 45):
        if str(ws.cell(r, 1).value or "").startswith("四、审计结论"):
            hint_cell = ws.cell(r + 1, 1)
            if not hint_cell.value:
                hint_cell.value = (
                    "结论参考：A.本期无合并范围外关联方工程物资购销交易；"
                    "B.关联交易真实、定价公允、披露恰当；"
                    "C.除下列异常事项外未见重大异常：______。"
                )
                hint_cell.font = Font(color="808080", size=9, italic=True)
            break


def main():
    updated = []
    for path in TARGETS:
        if not path.exists():
            print("SKIP missing", path)
            continue
        wb = openpyxl.load_workbook(path)
        if SHEET not in wb.sheetnames:
            print("SKIP no sheet", path)
            continue
        # idempotent: skip if already has purchase total formula
        ws = wb[SHEET]
        if ws["E17"].value == "=SUM(E11:E16)" and str(ws["A24"].value or "") == "合计":
            print("ALREADY", path)
            continue
        enhance(ws)
        wb.save(path)
        updated.append(str(path))
        print("UPDATED", path)
    print("done:", len(updated))


if __name__ == "__main__":
    main()

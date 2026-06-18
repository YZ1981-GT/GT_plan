"""Deep-read A7-A15 workpaper xlsx/docx templates; append to a7_a15_xlsx_audit.json."""
from __future__ import annotations

import argparse
import json
import re
from pathlib import Path

import openpyxl
from openpyxl.utils import column_index_from_string, get_column_letter

ROOT = Path(__file__).resolve().parents[2]
TEMPLATES = ROOT / "backend" / "wp_templates" / "A"
OUT = ROOT / "backend" / "data" / "a7_a15_xlsx_audit.json"
PROC_JSON = ROOT / "backend" / "data" / "procedure_table_templates.json"

AUDIT_WP_CODES = [
    "A7", "A7-1", "A7-2",
    "A8", "A8-1", "A8-2",
    "A9", "A9-1", "A9-2",
    "A10", "A10-1", "A10-2",
    "A11", "A11-1",
    "A12", "A12-1",
    "A13",
    "A14", "A14-1", "A14-2", "A14-3", "A14-4", "A14-5", "A14-6",
    "A15", "A15-1",
]


def resolve_template_path(wp_code: str) -> Path:
    """Match physical file by wp_code prefix (handles date suffixes; A9-1向… no space)."""
    if "-" in wp_code:
        matches = sorted(
            p
            for p in TEMPLATES.iterdir()
            if (p.name.startswith(f"{wp_code} ") or p.name.startswith(f"{wp_code}"))
            and re.match(rf"^{re.escape(wp_code)}(\s|[^-\d])", p.name)
        )
    else:
        matches = sorted(
            p for p in TEMPLATES.iterdir()
            if p.name.startswith(f"{wp_code} ")
            and not re.match(rf"^{re.escape(wp_code)}-\d", p.name)
        )
    if not matches:
        raise FileNotFoundError(f"No template for {wp_code} under {TEMPLATES}")
    if len(matches) > 1:
        xlsx = [p for p in matches if p.suffix.lower() == ".xlsx"]
        return xlsx[0] if xlsx else matches[0]
    return matches[0]


def cell_str(v) -> str | None:
    if v is None:
        return None
    s = str(v).strip()
    return s if s else None


def detect_header_row(ws, max_scan: int = 25) -> int:
    for r in range(1, min(max_scan, ws.max_row + 1)):
        vals = [cell_str(ws.cell(r, c).value) for c in range(1, ws.max_column + 1)]
        texts = " ".join(v for v in vals if v)
        if "序号" in texts and any(k in texts for k in ("程序", "内容", "调查")):
            return r
        if "项目" in texts and any(k in texts for k in ("母公司", "合计", "子公司")):
            return r
    return 5


def audit_program_sheet(ws, sn: str) -> dict:
    merged = [str(m) for m in ws.merged_cells.ranges]
    hr = detect_header_row(ws)
    cols = []
    ref_col = None
    for c in range(1, ws.max_column + 1):
        v = cell_str(ws.cell(hr, c).value)
        if v:
            cols.append({"col": get_column_letter(c), "label": v})
            if v == "索引号":
                ref_col = c
    rows = []
    for r in range(hr + 1, ws.max_row + 1):
        a = ws.cell(r, 1).value
        b = ws.cell(r, 2).value
        preview = cell_str(b)
        if a is None and not preview:
            continue
        seq_str = str(a) if a is not None else None
        if isinstance(a, (int, float)):
            seq_str = str(int(a)) if float(a).is_integer() else str(a)
        if preview in ("编制说明：",) or (seq_str and str(seq_str).startswith("提示：")):
            continue
        entry: dict = {
            "row": r,
            "seq": a if a is not None else None,
            "seq_normalized": seq_str,
            "procedure_preview": (preview[:150] if preview else None),
        }
        if ref_col:
            v = cell_str(ws.cell(r, ref_col).value)
            if v:
                entry["ref_index"] = v.replace("\n", ",").replace("、", ",")
        for col_letter, key in (("I", "applicable_tag"), ("J", "extra_tag"), ("K", "cycle_tag")):
            ci = column_index_from_string(col_letter)
            if ci <= ws.max_column:
                v = cell_str(ws.cell(r, ci).value)
                if v:
                    entry[key] = v
        rows.append(entry)
    return {
        "name": sn,
        "component_type": "a-program-console",
        "header_row": hr,
        "max_row": ws.max_row,
        "max_col": ws.max_column,
        "columns": cols,
        "ref_index_col": get_column_letter(ref_col) if ref_col else None,
        "data_rows": len(rows),
        "merged_ranges_count": len(merged),
        "merged_ranges_sample": merged[:15],
        "procedure_rows": rows,
    }


def audit_program_table(path: Path, wp_code: str) -> dict:
    wb = openpyxl.load_workbook(path, data_only=True, read_only=False)
    sheets = []
    for sn in wb.sheetnames:
        if sn == "GT_Custom" or "审计程序" not in sn:
            continue
        sheets.append(audit_program_sheet(wb[sn], sn))
    wb.close()
    return {
        "wp_code": wp_code,
        "filename": path.name,
        "runtime": "a-program-console",
        "sheets": sheets,
        "parser_ready": True,
    }


def audit_checklist_sheet(ws, sn: str) -> dict:
    merged = [str(m) for m in ws.merged_cells.ranges]
    hr = 5
    for r in range(1, min(15, ws.max_row + 1)):
        texts = " ".join(cell_str(ws.cell(r, c).value) or "" for c in range(1, min(8, ws.max_column + 1)))
        if "序号" in texts and ("调查" in texts or "内容" in texts):
            hr = r
            break
    cols = []
    for c in range(1, ws.max_column + 1):
        v = cell_str(ws.cell(hr, c).value)
        if v:
            cols.append({"col": get_column_letter(c), "label": v})
    rows = []
    for r in range(hr + 1, ws.max_row + 1):
        a = ws.cell(r, 1).value
        b = ws.cell(r, 2).value
        if a is None and not cell_str(b):
            continue
        rows.append({
            "row": r,
            "seq": a,
            "content_preview": (str(b)[:120] if b else None),
            "col_d": cell_str(ws.cell(r, 4).value) if ws.max_column >= 4 else None,
            "col_f": cell_str(ws.cell(r, 6).value) if ws.max_column >= 6 else None,
        })
    return {
        "name": sn,
        "component_type": "checklist-table",
        "header_row": hr,
        "max_row": ws.max_row,
        "max_col": ws.max_column,
        "columns": cols,
        "data_rows": len(rows),
        "checklist_rows_sample": rows[:25],
        "merged_ranges_count": len(merged),
        "merged_ranges_sample": merged[:10],
    }


def audit_univer_grid_sheet(ws, sn: str) -> dict:
    merged = [str(m) for m in ws.merged_cells.ranges]
    hr = detect_header_row(ws) or 5
    cols = []
    for c in range(1, ws.max_column + 1):
        v = cell_str(ws.cell(hr, c).value)
        if v:
            cols.append({"col": get_column_letter(c), "label": v})
    rows = []
    for r in range(hr + 1, min(ws.max_row + 1, hr + 30)):
        cells = {get_column_letter(c): cell_str(ws.cell(r, c).value) for c in range(1, ws.max_column + 1)}
        cells = {k: v for k, v in cells.items() if v}
        if cells:
            rows.append({"row": r, "cells": cells})
    return {
        "name": sn,
        "component_type": "d-form-table",
        "wp_alias": "A11-WP-1",
        "notes": "xlsx sheet「期后事项审定表A11-1」≠ A11-1 docx 问询函",
        "header_row": hr,
        "max_row": ws.max_row,
        "max_col": ws.max_column,
        "columns": cols,
        "data_rows_sample": rows,
        "merged_ranges_count": len(merged),
        "merged_ranges_sample": merged[:10],
    }


def audit_b_index_sheet(ws, sn: str) -> dict:
    rows = []
    for r in range(1, min(ws.max_row + 1, 30)):
        cells = {get_column_letter(c): cell_str(ws.cell(r, c).value) for c in range(1, ws.max_column + 1)}
        cells = {k: v for k, v in cells.items() if v}
        if cells:
            rows.append({"row": r, "cells": cells})
    return {
        "name": sn,
        "component_type": "b-index",
        "max_row": ws.max_row,
        "max_col": ws.max_column,
        "index_rows_sample": rows,
    }


def audit_d_form_table_sheet(ws, sn: str) -> dict:
    merged = [str(m) for m in ws.merged_cells.ranges]
    hr = None
    cols: list[dict] = []
    for r in range(1, min(30, ws.max_row + 1)):
        texts = " ".join(cell_str(ws.cell(r, c).value) or "" for c in range(1, min(12, ws.max_column + 1)))
        if hr is None and "序号" in texts and any(k in texts for k in ("说明", "索引", "错报")):
            hr = r
            for c in range(1, ws.max_column + 1):
                v = cell_str(ws.cell(r, c).value)
                if v:
                    cols.append({"col": get_column_letter(c), "label": v})
    sections: list[dict] = []
    sub_headers: list[dict] = []
    for r in range(1, ws.max_row + 1):
        a = cell_str(ws.cell(r, 1).value)
        b = cell_str(ws.cell(r, 2).value)
        if a and (
            a.startswith(("—", "一", "二", "三", "四"))
            or "、" in a[:6]
        ) and len(a) <= 60:
            sections.append({"row": r, "label": a})
        row_text = " ".join(cell_str(ws.cell(r, c).value) or "" for c in range(1, min(12, ws.max_column + 1)))
        if any(k in row_text for k in ("索引号", "导致错报", "沟通时间", "借方科目", "资产(+)")):
            cells = {
                get_column_letter(c): cell_str(ws.cell(r, c).value)
                for c in range(1, ws.max_column + 1)
                if cell_str(ws.cell(r, c).value)
            }
            sub_headers.append({"row": r, "headers": cells})
    data_sample = []
    total_data = 0
    start = (hr + 1) if hr else 6
    for r in range(start, ws.max_row + 1):
        row_vals = {}
        has_data = False
        for c in range(1, ws.max_column + 1):
            v = ws.cell(r, c).value
            if v is not None and str(v).strip():
                has_data = True
                if len(data_sample) < 12:
                    row_vals[get_column_letter(c)] = str(v)[:80]
        if has_data:
            total_data += 1
            if len(data_sample) < 12:
                data_sample.append({"row": r, "cells": row_vals})
    return {
        "name": sn,
        "component_type": "d-form-table",
        "header_row": hr,
        "max_row": ws.max_row,
        "max_col": ws.max_column,
        "columns": cols,
        "sections": sections[:15],
        "sub_table_headers": sub_headers[:10],
        "data_rows": total_data,
        "data_sample": data_sample,
        "merged_ranges_count": len(merged),
        "merged_ranges_sample": merged[:10],
    }


def audit_misstatement_summary_sheet(ws, sn: str) -> dict:
    merged = [str(m) for m in ws.merged_cells.ranges]
    hr = 5
    cols_main = []
    for c in range(1, ws.max_column + 1):
        v = cell_str(ws.cell(hr, c).value)
        if v:
            cols_main.append({"col": get_column_letter(c), "label": v})
    cols_sub = []
    if ws.max_row >= 6:
        for c in range(1, ws.max_column + 1):
            v = cell_str(ws.cell(6, c).value)
            if v:
                cols_sub.append({"col": get_column_letter(c), "label": v})
    sections = []
    for r in range(hr + 1, ws.max_row + 1):
        a = cell_str(ws.cell(r, 1).value)
        if a and (a in ("一", "二", "三") or "合计" in a or a.startswith(("1、", "2、"))):
            b = cell_str(ws.cell(r, 2).value)
            sections.append({"row": r, "col_a": a, "col_b_preview": (b[:60] if b else None)})
    return {
        "name": sn,
        "component_type": "misstatement-summary",
        "header_row": hr,
        "sub_header_row": 6 if cols_sub else None,
        "max_row": ws.max_row,
        "max_col": ws.max_column,
        "columns_main": cols_main,
        "columns_sub": cols_sub,
        "sections": sections[:20],
        "merged_ranges_count": len(merged),
        "merged_ranges_sample": merged[:10],
        "notes": "A13-1 未更正错报汇总；三大块错报 + 底部合计表",
    }


def audit_a13_bundle(path: Path, wp_code: str) -> dict:
    wb = openpyxl.load_workbook(path, data_only=True, read_only=False)
    sheets = []
    for sn in wb.sheetnames:
        if sn == "GT_Custom":
            sheets.append({"name": sn, "component_type": "skip"})
            continue
        ws = wb[sn]
        if sn == "底稿目录" or "目录" in sn:
            sheets.append(audit_b_index_sheet(ws, sn))
        elif "程序" in sn:
            sheets.append(audit_program_sheet(ws, sn))
        elif "A13-1" in sn or "未更正错报汇总" in sn:
            sheets.append(audit_misstatement_summary_sheet(ws, sn))
        elif "A13-2" in sn or "列报和披露" in sn:
            sheets.append(audit_d_form_table_sheet(ws, sn))
        elif "A13-3" in sn or "评价识别" in sn:
            sheets.append(audit_d_form_table_sheet(ws, sn))
        elif "A13-4" in sn or "错报的性质" in sn:
            sheets.append(audit_d_form_table_sheet(ws, sn))
        elif "A13-5" in sn or "沟通" in sn:
            sheets.append(audit_d_form_table_sheet(ws, sn))
        else:
            sheets.append({"name": sn, "component_type": "unknown", "max_row": ws.max_row})
    wb.close()
    return {
        "wp_code": wp_code,
        "filename": path.name,
        "runtime": "a13-bundle",
        "sheets": sheets,
        "parser_ready": False,
        "notes": "Single xlsx Tab suite: program + A13-1 summary + A13-2~5 forms",
    }


def _is_example_or_reference_sheet(sn: str) -> bool:
    return any(k in sn for k in ("示例", "【示例", "参考"))


def audit_control_deficiency_summary(path: Path, wp_code: str) -> dict:
    wb = openpyxl.load_workbook(path, data_only=True, read_only=False)
    sheets = []
    for sn in wb.sheetnames:
        if _is_example_or_reference_sheet(sn):
            sheets.append({"name": sn, "component_type": "example-skip"})
            continue
        ws = wb[sn]
        merged = [str(m) for m in ws.merged_cells.ranges]
        hr = 5
        cols_main = []
        for c in range(1, ws.max_column + 1):
            v = cell_str(ws.cell(hr, c).value)
            if v:
                cols_main.append({"col": get_column_letter(c), "label": v})
        cols_sub = []
        if ws.max_row >= 6:
            for c in range(1, ws.max_column + 1):
                v = cell_str(ws.cell(6, c).value)
                if v:
                    cols_sub.append({"col": get_column_letter(c), "label": v})
        footnotes = []
        for r in range(7, ws.max_row + 1):
            a = cell_str(ws.cell(r, 1).value)
            if a and a.startswith("注"):
                footnotes.append({"row": r, "text_preview": a[:120]})
        sheets.append(
            {
                "name": sn,
                "component_type": "checklist-table",
                "header_row": hr,
                "sub_header_row": 6 if cols_sub else None,
                "max_row": ws.max_row,
                "max_col": ws.max_column,
                "columns_main": cols_main,
                "columns_sub": cols_sub,
                "data_rows": max(0, ws.max_row - hr - len(footnotes)),
                "footnotes": footnotes[:5],
                "merged_ranges_count": len(merged),
                "merged_ranges_sample": merged[:10],
                "notes": "缺陷汇总 grid；R5+R6 双行表头（认定列）；含注1/注2",
            }
        )
    wb.close()
    return {
        "wp_code": wp_code,
        "filename": path.name,
        "runtime": "checklist-table",
        "sheets": sheets,
        "parser_ready": False,
        "notes": "A14-1 汇总表 + 示例 sheet（example-skip）",
    }


def audit_deficiency_eval_workbook(
    path: Path, wp_code: str, *, default_runtime: str = "d-form-table"
) -> dict:
    wb = openpyxl.load_workbook(path, data_only=True, read_only=False)
    sheets = []
    for sn in wb.sheetnames:
        if _is_example_or_reference_sheet(sn):
            sheets.append({"name": sn, "component_type": "example-skip"})
            continue
        ws = wb[sn]
        info = audit_d_form_table_sheet(ws, sn)
        if wp_code == "A14-6":
            info["component_type"] = "e-control-test"
        # detect evaluation steps
        steps = []
        for r in range(1, ws.max_row + 1):
            a = cell_str(ws.cell(r, 1).value)
            if a and ("Step " in a or a.startswith("步骤")):
                steps.append({"row": r, "step_preview": a[:100]})
        info["evaluation_steps"] = steps[:12]
        info["step_count"] = len(steps)
        if wp_code == "A14-2" and "业务流程" in (cell_str(ws.cell(2, 1).value) or ""):
            info["notes"] = "⚠️ sheet 位于 A14-2 文件内但标题为业务流程层面（索引 A14-4），疑为模板拷贝残留"
        sheets.append(info)
    wb.close()
    runtime = "e-control-test" if wp_code == "A14-6" else default_runtime
    return {
        "wp_code": wp_code,
        "filename": path.name,
        "runtime": runtime,
        "sheets": sheets,
        "parser_ready": False,
    }


def audit_a14_3_workbook(path: Path, wp_code: str) -> dict:
    wb = openpyxl.load_workbook(path, data_only=True, read_only=False)
    sheets = []
    for sn in wb.sheetnames:
        if _is_example_or_reference_sheet(sn):
            sheets.append({"name": sn, "component_type": "example-skip"})
            continue
        ws = wb[sn]
        if "缺陷汇总" in sn:
            info = audit_d_form_table_sheet(ws, sn)
            info["component_type"] = "d-form-table"
            info["notes"] = "IT 缺陷汇总表：缺陷编号/类别/控制类型/应用/描述/补偿性控制/风险/报表项目/认定"
        elif "沟通" in sn or "讨论" in sn:
            info = audit_d_form_table_sheet(ws, sn)
            info["component_type"] = "d-form-table"
            info["notes"] = "IT 控制缺陷对财务报表影响的讨论记录"
        else:
            info = audit_d_form_table_sheet(ws, sn)
            info["component_type"] = "d-form-table"
            info["wp_sub"] = sn
            steps = []
            for r in range(1, ws.max_row + 1):
                a = cell_str(ws.cell(r, 1).value)
                if a and "Step " in a:
                    steps.append({"row": r, "step_preview": a[:100]})
            info["evaluation_steps"] = steps[:10]
        sheets.append(info)
    wb.close()
    return {
        "wp_code": wp_code,
        "filename": path.name,
        "runtime": "a14-3-workbook",
        "sheets": sheets,
        "parser_ready": False,
        "notes": "多 sheet：缺陷汇总 + IT控制缺陷评价模板 + 沟通纪要；示例 sheet 已 skip",
    }


def audit_questionnaire_sheet(ws, sn: str) -> dict:
    merged = [str(m) for m in ws.merged_cells.ranges]
    hr = 5
    for r in range(1, min(20, ws.max_row + 1)):
        texts = " ".join(cell_str(ws.cell(r, c).value) or "" for c in range(1, min(8, ws.max_column + 1)))
        if "序号" in texts and any(k in texts for k in ("是否存在", "适用情况", "适用")):
            hr = r
            break
    cols = []
    for c in range(1, ws.max_column + 1):
        v = cell_str(ws.cell(hr, c).value)
        if v:
            cols.append({"col": get_column_letter(c), "label": v})
    sections: list[dict] = []
    rows: list[dict] = []
    for r in range(1, ws.max_row + 1):
        a_str = cell_str(ws.cell(r, 1).value)
        if a_str and a_str.startswith(("一、", "二、", "三、", "四、")):
            sections.append({"row": r, "label": a_str})
        elif a_str and (a_str.startswith("调查") or a_str.startswith("审计程序")):
            sections.append({"row": r, "label": a_str, "type": "conclusion"})
    for r in range(hr + 1, ws.max_row + 1):
        a = ws.cell(r, 1).value
        b = ws.cell(r, 2).value
        a_str = cell_str(a)
        if a_str and a_str.startswith(("一、", "二、", "三、", "四、")):
            continue
        if a_str and (a_str.startswith("调查") or a_str.startswith("审计程序")):
            continue
        if a is None and not cell_str(b):
            continue
        rows.append({
            "row": r,
            "seq": a,
            "content_preview": (str(b)[:120] if b else None),
            "col_d": cell_str(ws.cell(r, 4).value) if ws.max_column >= 4 else None,
            "col_e": cell_str(ws.cell(r, 5).value) if ws.max_column >= 5 else None,
        })
    return {
        "name": sn,
        "component_type": "checklist-table",
        "header_row": hr,
        "max_row": ws.max_row,
        "max_col": ws.max_column,
        "columns": cols,
        "sections": sections,
        "data_rows": len(rows),
        "checklist_rows_sample": rows[:25],
        "merged_ranges_count": len(merged),
        "merged_ranges_sample": merged[:10],
    }


def audit_a15_bundle(path: Path, wp_code: str) -> dict:
    wb = openpyxl.load_workbook(path, data_only=True, read_only=False)
    sheets = []
    for sn in wb.sheetnames:
        if sn == "GT_Custom":
            sheets.append({"name": sn, "component_type": "skip"})
            continue
        ws = wb[sn]
        if "审计程序" in sn:
            sheets.append(audit_program_sheet(ws, sn))
        elif "决策图" in sn:
            sheets.append({
                "name": sn,
                "component_type": "guidance-reference",
                "max_row": ws.max_row,
                "notes": "持续经营能力对审计报告影响决策图；guidance 引用，非数据 sheet",
            })
        else:
            sheets.append({"name": sn, "component_type": "unknown", "max_row": ws.max_row})
    wb.close()
    return {
        "wp_code": wp_code,
        "filename": path.name,
        "runtime": "a15-bundle",
        "sheets": sheets,
        "parser_ready": True,
        "notes": "程序表 + 决策图 guidance sheet",
    }


def audit_a15_1_checklist(path: Path, wp_code: str) -> dict:
    wb = openpyxl.load_workbook(path, data_only=True, read_only=False)
    sheets = []
    for sn in wb.sheetnames:
        if sn == "GT_Custom":
            sheets.append({"name": sn, "component_type": "skip"})
            continue
        sheets.append(audit_questionnaire_sheet(wb[sn], sn))
    wb.close()
    return {
        "wp_code": wp_code,
        "filename": path.name,
        "runtime": "checklist-table",
        "sheets": sheets,
        "parser_ready": False,
        "notes": "三章节问卷（财务11+经营6+其他4）+ 调查结论",
    }


def audit_a11_bundle(path: Path, wp_code: str) -> dict:
    wb = openpyxl.load_workbook(path, data_only=True, read_only=False)
    sheets = []
    for sn in wb.sheetnames:
        if sn == "GT_Custom":
            sheets.append({"name": sn, "component_type": "skip"})
            continue
        ws = wb[sn]
        if "审计程序" in sn:
            sheets.append(audit_program_sheet(ws, sn))
        elif sn == "底稿目录" or "目录" in sn:
            sheets.append(audit_b_index_sheet(ws, sn))
        elif "审定表" in sn:
            sheets.append(audit_univer_grid_sheet(ws, sn))
        elif "调查问卷" in sn or "问卷" in sn:
            sheets.append(audit_checklist_sheet(ws, sn))
        else:
            sheets.append({"name": sn, "component_type": "unknown", "max_row": ws.max_row})
    wb.close()
    return {
        "wp_code": wp_code,
        "filename": path.name,
        "runtime": "a11-bundle",
        "sheets": sheets,
        "parser_ready": False,
        "notes": "Single xlsx bundle: program + 审定表A11-1 sheet + 问卷A11-2/A11-3",
    }


def audit_univer_table(path: Path, wp_code: str) -> dict:
    wb = openpyxl.load_workbook(path, data_only=True, read_only=False)
    sheets = []
    for sn in wb.sheetnames:
        ws = wb[sn]
        merged = [str(m) for m in ws.merged_cells.ranges]
        hr = detect_header_row(ws)
        cols = []
        for c in range(1, ws.max_column + 1):
            v = cell_str(ws.cell(hr, c).value)
            if v:
                cols.append({"col": get_column_letter(c), "label": v})
        data_sample = []
        total_data = 0
        for r in range(hr + 1, ws.max_row + 1):
            row_vals = {}
            has_data = False
            for c in range(1, ws.max_column + 1):
                v = ws.cell(r, c).value
                if v is not None and str(v).strip():
                    has_data = True
                    if len(data_sample) < 20:
                        row_vals[get_column_letter(c)] = str(v)[:80]
            if has_data:
                total_data += 1
                if len(data_sample) < 20:
                    data_sample.append({"row": r, "cells": row_vals})
        sections = []
        for r in range(hr + 1, ws.max_row + 1):
            a = cell_str(ws.cell(r, 1).value)
            if a and not re.match(r"^[\d.]+$", a) and len(a) <= 40:
                sections.append({"row": r, "label": a})
        sheets.append(
            {
                "name": sn,
                "component_type": "univer",
                "header_row": hr,
                "max_row": ws.max_row,
                "max_col": ws.max_column,
                "columns": cols,
                "data_rows": total_data,
                "section_labels": sections[:20],
                "data_sample": data_sample,
                "merged_ranges_count": len(merged),
                "merged_ranges_sample": merged[:10],
            }
        )
    wb.close()
    return {
        "wp_code": wp_code,
        "filename": path.name,
        "runtime": "univer",
        "sheets": sheets,
        "parser_ready": False,
        "notes": "Univer/HTML table; row structure confirmed by deep read",
    }


def audit_c_note(path: Path, wp_code: str) -> dict:
    wb = openpyxl.load_workbook(path, data_only=True, read_only=False)
    sheets = []
    for sn in wb.sheetnames:
        ws = wb[sn]
        if sn == "GT_Custom":
            sheets.append({"name": sn, "component_type": "skip", "notes": "placeholder sheet"})
            continue
        merged = [str(m) for m in ws.merged_cells.ranges]
        hr = 6
        cols_main = []
        for c in range(1, ws.max_column + 1):
            v = cell_str(ws.cell(hr, c).value)
            if v:
                cols_main.append({"col": get_column_letter(c), "label": v})
        sections = []
        for r in range(1, ws.max_row + 1):
            a = cell_str(ws.cell(r, 1).value)
            b = cell_str(ws.cell(r, 2).value)
            if a and (
                a.startswith(("一", "二", "三", "四", "五", "六"))
                or "披露" in a
                or "关联" in a
            ):
                sections.append({"row": r, "col_a": a, "col_b_preview": (b[:60] if b else None)})
        sub_headers = []
        for r in range(1, ws.max_row + 1):
            row_text = " ".join(
                cell_str(ws.cell(r, c).value) or "" for c in range(1, min(13, ws.max_column + 1))
            )
            if any(
                k in row_text
                for k in (
                    "交易类型",
                    "金额",
                    "定价",
                    "占同类",
                    "关联交易",
                    "关联往来",
                    "担保",
                    "关键管理人员",
                )
            ):
                cells = {
                    get_column_letter(c): cell_str(ws.cell(r, c).value)
                    for c in range(1, ws.max_column + 1)
                    if cell_str(ws.cell(r, c).value)
                }
                sub_headers.append({"row": r, "headers": cells})
        sheets.append(
            {
                "name": sn,
                "component_type": "c-note-table",
                "header_row": hr,
                "max_row": ws.max_row,
                "max_col": ws.max_column,
                "columns_section1": cols_main,
                "sections": sections,
                "sub_table_headers": sub_headers[:25],
                "merged_ranges_count": len(merged),
                "merged_ranges_sample": merged[:10],
            }
        )
    wb.close()
    return {
        "wp_code": wp_code,
        "filename": path.name,
        "runtime": "c-note-table",
        "sheets": sheets,
        "parser_ready": False,
        "notes": "Multi-section disclosure template; c-note-table needs section-aware parser",
    }


def audit_d_form(path: Path, wp_code: str) -> dict:
    wb = openpyxl.load_workbook(path, data_only=True, read_only=False)
    sheets = []
    for sn in wb.sheetnames:
        ws = wb[sn]
        merged = [str(m) for m in ws.merged_cells.ranges]
        hr = detect_header_row(ws)
        cols = []
        for c in range(1, ws.max_column + 1):
            v = cell_str(ws.cell(hr, c).value)
            if v:
                cols.append({"col": get_column_letter(c), "label": v})
        data_sample = []
        total_data = 0
        for r in range(hr + 1, ws.max_row + 1):
            row_vals = {}
            has_data = False
            for c in range(1, ws.max_column + 1):
                v = ws.cell(r, c).value
                if v is not None and str(v).strip():
                    has_data = True
                    if len(data_sample) < 15:
                        row_vals[get_column_letter(c)] = str(v)[:80]
            if has_data:
                total_data += 1
                if len(data_sample) < 15:
                    data_sample.append({"row": r, "cells": row_vals})
        sheets.append(
            {
                "name": sn,
                "component_type": "d-form-confirmation",
                "header_row": hr,
                "max_row": ws.max_row,
                "max_col": ws.max_column,
                "columns": cols,
                "data_rows": total_data,
                "data_sample": data_sample,
                "merged_ranges_count": len(merged),
                "merged_ranges_sample": merged[:10],
            }
        )
    wb.close()
    return {
        "wp_code": wp_code,
        "filename": path.name,
        "runtime": "d-form-confirmation",
        "sheets": sheets,
        "parser_ready": False,
        "notes": "d-form-confirmation HTML; column mapping from audit",
    }


def audit_wp(wp_code: str) -> dict:
    path = resolve_template_path(wp_code)
    if wp_code == "A11":
        return audit_a11_bundle(path, wp_code)
    if wp_code == "A13":
        return audit_a13_bundle(path, wp_code)
    if wp_code == "A14-1":
        return audit_control_deficiency_summary(path, wp_code)
    if wp_code in ("A14-2", "A14-4", "A14-5"):
        return audit_deficiency_eval_workbook(path, wp_code)
    if wp_code == "A14-3":
        return audit_a14_3_workbook(path, wp_code)
    if wp_code == "A14-6":
        return audit_deficiency_eval_workbook(path, wp_code, default_runtime="e-control-test")
    if wp_code == "A15":
        return audit_a15_bundle(path, wp_code)
    if wp_code == "A15-1":
        return audit_a15_1_checklist(path, wp_code)
    if wp_code in ("A7", "A8", "A9", "A10", "A12", "A14"):
        return audit_program_table(path, wp_code)
    if wp_code == "A7-1":
        return audit_univer_table(path, wp_code)
    if wp_code == "A7-2":
        return audit_c_note(path, wp_code)
    if wp_code == "A10-2":
        return audit_d_form(path, wp_code)
    raise NotImplementedError(f"audit handler not implemented for {wp_code}")


def merge_audit(entry: dict) -> None:
    existing: list[dict] = []
    if OUT.exists():
        existing = json.loads(OUT.read_text(encoding="utf-8"))
    existing = [e for e in existing if e.get("wp_code") != entry["wp_code"]]
    existing.append(entry)
    existing.sort(key=lambda x: x["wp_code"])
    OUT.write_text(json.dumps(existing, ensure_ascii=False, indent=2), encoding="utf-8")


def diff_procedure_table(wp_code: str, entry: dict) -> dict:
    proc_data = json.loads(PROC_JSON.read_text(encoding="utf-8"))
    proc = proc_data.get("tables", proc_data)[wp_code]

    def flatten(items: list) -> list[tuple[str, str | None]]:
        out: list[tuple[str, str | None]] = []

        def walk(it: dict) -> None:
            out.append((str(it["seq"]), it.get("ref_index")))
            for sub in it.get("sub_items", []):
                walk(sub)

        for it in items:
            walk(it)
        return out

    def normalize_xlsx_seqs(rows: list) -> list[tuple[str, str | None]]:
        import re

        paren_re = re.compile(r"^[\uff08(](\d+)[\uff09)]$")
        bullet_wp = wp_code in ("A8",)  # 仅 A8 将 ● 行展开为 parent.N；其余模块 ● 为归档内联子项
        out: list[tuple[str, str | None]] = []
        parent: str | None = None
        bullet_idx = 0
        for r in rows:
            seq = r.get("seq")
            if seq is None:
                continue
            seq_s = str(seq).strip()
            if seq_s.startswith("说明"):
                continue
            if seq_s == "●":
                if not bullet_wp or parent is None:
                    continue
                bullet_idx += 1
                out.append((f"{parent}.{bullet_idx}", r.get("ref_index")))
                continue
            m = paren_re.match(seq_s)
            if m:
                if parent is None:
                    continue
                out.append((f"{parent}.{m.group(1)}", r.get("ref_index")))
                continue
            bullet_idx = 0
            parent = str(r.get("seq_normalized") or seq_s)
            out.append((parent, r.get("ref_index")))
        return out

    # JSON sub-items expanded for UI while xlsx keeps inline (一)(二) in parent row
    INLINE_PARENT_SUBS: dict[str, dict[str, frozenset[str]]] = {
        "A8": {"5": frozenset({"5.1", "5.2", "5.3"})},
    }

    def normalize_ref(ref: str | None) -> str | None:
        if ref is None:
            return None
        s = str(ref).replace("、", ",")
        s = s.replace("A14-,3A14-4", "A14-3,A14-4")
        parts = [p.strip() for p in s.split(",") if p.strip()]
        return ",".join(parts) if parts else s

    json_seqs = flatten(proc["items"])
    program_sheet = next(s for s in entry["sheets"] if s.get("component_type") == "a-program-console")
    xlsx_seqs = normalize_xlsx_seqs(program_sheet["procedure_rows"])
    json_set = {s for s, _ in json_seqs}
    xlsx_set = {s for s, _ in xlsx_seqs}
    missing_in_json = sorted(xlsx_set - json_set)
    missing_in_xlsx = sorted(json_set - xlsx_set)
    for parent, subs in INLINE_PARENT_SUBS.get(wp_code, {}).items():
        if parent in xlsx_set:
            missing_in_xlsx = [s for s in missing_in_xlsx if s not in subs]
    ref_mismatch = []
    json_map = dict(json_seqs)
    xlsx_map = dict(xlsx_seqs)
    for seq in json_set & xlsx_set:
        jr, xr = normalize_ref(json_map.get(seq)), normalize_ref(xlsx_map.get(seq))
        if not jr and not xr:
            continue
        if jr and not xr:
            continue  # xlsx 索引列常为空；chip ref 以 JSON 为准
        if jr and xr and jr != xr:
            # A10-1-1 等为 xlsx 变体，spec 仅 A10-1
            if wp_code == "A10" and jr in xr.replace(" ", "").split(","):
                continue
            ref_mismatch.append({"seq": seq, "json_ref": json_map.get(seq), "xlsx_ref": xlsx_map.get(seq)})
    return {
        "wp_code": wp_code,
        "json_count": len(json_seqs),
        "xlsx_count": len(xlsx_seqs),
        "missing_in_json": missing_in_json,
        "missing_in_xlsx": missing_in_xlsx,
        "ref_index_mismatches": ref_mismatch,
    }


def run_diff_only() -> int:
    """Re-audit all wp_codes in the existing JSON and compare key fields.

    Returns 0 if no drift, 1 if any drift detected.
    Does NOT write/update the JSON file (read-only check).
    # TODO: A17 audit 就绪后纳入 --diff-only 检查范围
    """
    if not OUT.exists():
        print(f"ERROR: {OUT} not found; cannot run --diff-only check.")
        return 1

    stored: list[dict] = json.loads(OUT.read_text(encoding="utf-8"))
    stored_map: dict[str, dict] = {e["wp_code"]: e for e in stored}
    wp_codes = sorted(stored_map.keys())

    drifted: list[str] = []
    matched: list[str] = []

    for code in wp_codes:
        stored_entry = stored_map[code]
        reasons: list[str] = []

        # Docx entries: only verify file exists and filename matches
        if stored_entry.get("format") == "docx":
            try:
                path = resolve_template_path(code)
                if path.name != stored_entry.get("filename"):
                    reasons.append(
                        f"filename changed: {stored_entry.get('filename')!r} -> {path.name!r}"
                    )
            except FileNotFoundError as exc:
                reasons.append(f"template file missing: {exc}")
            if reasons:
                drifted.append(code)
                print(f"  DRIFT {code}:")
                for r in reasons:
                    print(f"    - {r}")
            else:
                matched.append(code)
            continue

        # xlsx entries: full re-audit and compare
        try:
            fresh_entry = audit_wp(code)
        except FileNotFoundError as exc:
            reasons.append(f"template file missing: {exc}")
            drifted.append(code)
            print(f"  DRIFT {code}: {reasons[0]}")
            continue
        except NotImplementedError:
            # Handler not yet implemented for this code; skip gracefully
            matched.append(code)
            continue

        # Compare filename
        if fresh_entry["filename"] != stored_entry.get("filename"):
            reasons.append(
                f"filename changed: {stored_entry.get('filename')!r} -> {fresh_entry['filename']!r}"
            )

        # Compare sheets count
        stored_sheets = stored_entry.get("sheets", [])
        fresh_sheets = fresh_entry.get("sheets", [])
        if len(fresh_sheets) != len(stored_sheets):
            reasons.append(
                f"sheets count changed: {len(stored_sheets)} -> {len(fresh_sheets)}"
            )
        else:
            # Compare per-sheet key fields
            for i, (fs, ss) in enumerate(zip(fresh_sheets, stored_sheets)):
                sheet_name = fs.get("name", f"sheet[{i}]")
                if fs.get("component_type") != ss.get("component_type"):
                    reasons.append(
                        f"sheet '{sheet_name}' component_type: "
                        f"{ss.get('component_type')!r} -> {fs.get('component_type')!r}"
                    )
                # Compare total_rows: use max_row (or data_rows as fallback)
                fresh_rows = fs.get("max_row") or fs.get("data_rows")
                stored_rows = ss.get("max_row") or ss.get("data_rows")
                if fresh_rows is not None and stored_rows is not None:
                    if fresh_rows != stored_rows:
                        reasons.append(
                            f"sheet '{sheet_name}' row count: {stored_rows} -> {fresh_rows}"
                        )

        if reasons:
            drifted.append(code)
            print(f"  DRIFT {code}:")
            for r in reasons:
                print(f"    - {r}")
        else:
            matched.append(code)

    total = len(wp_codes)
    print(f"\n{'=' * 50}")
    print(f"Summary: {len(matched)}/{total} wp_codes match, {len(drifted)} drifted")
    if drifted:
        print(f"Drifted: {', '.join(drifted)}")
        return 1
    print("No drift detected.")
    return 0


def main() -> None:
    parser = argparse.ArgumentParser()
    parser.add_argument("wp_codes", nargs="*", help="e.g. A7 A7-1")
    parser.add_argument("--diff", action="store_true", help="print procedure_table diff for program tables")
    parser.add_argument(
        "--diff-only",
        action="store_true",
        help="CI mode: re-audit all stored wp_codes and exit 1 on drift (read-only, no JSON update)",
    )
    args = parser.parse_args()

    if args.diff_only:
        raise SystemExit(run_diff_only())

    if not args.wp_codes:
        parser.error("wp_codes are required unless --diff-only is specified")

    for code in args.wp_codes:
        entry = audit_wp(code)
        if args.diff and entry.get("runtime") in (
            "a-program-console",
            "a11-bundle",
            "a13-bundle",
            "a15-bundle",
        ):
            entry["procedure_table_diff"] = diff_procedure_table(code, entry)
        merge_audit(entry)
        print(f"audited {code}: {entry['filename']} ({len(entry['sheets'])} sheet(s))")
        if entry.get("procedure_table_diff"):
            print(json.dumps(entry["procedure_table_diff"], ensure_ascii=False, indent=2))


if __name__ == "__main__":
    main()

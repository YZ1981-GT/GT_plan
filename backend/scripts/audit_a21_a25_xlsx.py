"""Deep-read A21~A25 各角色复核 xlsx → a21_a25_xlsx_audit.json + review_definitions.

Tasks X-A21-1 ~ X-A25-2（a21-a25-review-workpapers spec audit 阶段）

用法:
    python scripts/audit_a21_a25_xlsx.py
    python scripts/audit_a21_a25_xlsx.py A21-1
    python scripts/audit_a21_a25_xlsx.py --write-definitions
"""
from __future__ import annotations

import argparse
import json
import re
from pathlib import Path

import openpyxl
from openpyxl.utils import get_column_letter

ROOT = Path(__file__).resolve().parents[2]
TEMPLATES = ROOT / "backend" / "wp_templates" / "A"
OUT_AUDIT = ROOT / "backend" / "data" / "a21_a25_xlsx_audit.json"
OUT_DEFS = ROOT / "backend" / "data" / "a21_a25_review_definitions.json"

# (wp_code, enterprise_variant) — variant None 或 large_soe / non_soe
AUDIT_TARGETS: list[tuple[str, str | None]] = [
    ("A21-1", None),
    ("A21-2", None),
    ("A22-1", None),
    ("A22-2", None),
    ("A23-1", None),
    ("A23-2", None),
    ("A24-1", "large_soe"),
    ("A24-1", "non_soe"),
    ("A24-2", None),
    ("A25-1", "large_soe"),
    ("A25-1", "non_soe"),
    ("A25-2", None),
]

ITEM_RE = re.compile(r"^[（(](\d+)[）)]")
SEQ_COL_RE = re.compile(r"^[\d]+(?:\.[\d]+)*$")
SKIP_SHEETS = {"GT_Custom", "底稿目录"}

ROLE_META: dict[str, dict] = {
    "A21": {
        "role": "field_lead",
        "role_label": "项目现场负责人",
        "applicable_categories": ["A", "B", "C"],
    },
    "A22": {
        "role": "manager",
        "role_label": "项目负责经理",
        "applicable_categories": ["A", "B", "C"],
    },
    "A23": {
        "role": "partner",
        "role_label": "项目合伙人",
        "applicable_categories": ["A", "B"],
    },
    "A24": {
        "role": "quality_reviewer",
        "role_label": "质量复核合伙人",
        "applicable_categories": ["A"],
    },
    "A25": {
        "role": "eqcr",
        "role_label": "质量控制复核人",
        "applicable_categories": ["A"],
    },
}


def cell_str(v) -> str | None:
    if v is None:
        return None
    s = str(v).strip()
    return s if s else None


def resolve_template_path(wp_code: str, enterprise_variant: str | None) -> Path:
    """按 wp_code + 国企变体匹配物理文件。"""
    candidates = [
        p
        for p in TEMPLATES.iterdir()
        if p.suffix.lower() == ".xlsx" and p.name.replace("  ", " ").startswith(f"{wp_code} ")
    ]
    if not candidates:
        raise FileNotFoundError(f"No xlsx for {wp_code} under {TEMPLATES}")

    if enterprise_variant == "large_soe":
        matched = [
            p for p in candidates
            if "大型国企" in p.name and "非大型" not in p.name
        ]
    elif enterprise_variant == "non_soe":
        matched = [
            p for p in candidates
            if "非大型国企" in p.name or "适用非大型" in p.name
        ]
    else:
        matched = [p for p in candidates if "大型国企" not in p.name and "非大型国企" not in p.name]

    if not matched:
        matched = candidates
    return sorted(matched)[0]


def _find_checklist_sheet(wb) -> tuple[str, int]:
    """返回 (sheet_name, header_row) — 表头含 是/否。"""
    best: tuple[str, int] | None = None
    for sn in wb.sheetnames:
        if sn in SKIP_SHEETS:
            continue
        if sn == "复核记录" or (sn.endswith("记录") and "复核" in sn and "表" not in sn):
            continue
        ws = wb[sn]
        for r in range(1, min(15, ws.max_row + 1)):
            texts = [
                cell_str(ws.cell(r, c).value) or ""
                for c in range(1, min(ws.max_column + 1, 15))
            ]
            if "是" in texts and "否" in texts:
                if best is None or len(sn) > len(best[0]):
                    best = (sn, r)
    if best is None:
        raise ValueError("checklist sheet not found")
    return best


def _find_record_sheet(wb) -> str | None:
    for sn in wb.sheetnames:
        if "复核记录" in sn or sn.strip() == "复核记录":
            return sn
    return None


def _detect_column_roles(ws, header_row: int) -> dict[str, str]:
    roles: dict[str, str] = {}
    for c in range(1, ws.max_column + 1):
        v = cell_str(ws.cell(header_row, c).value)
        if not v:
            continue
        letter = get_column_letter(c)
        if v == "是":
            roles["yes"] = letter
        elif v == "否":
            roles["no"] = letter
        elif "不适用" in v:
            roles["na"] = letter
        elif "复核记录" in v or "索引" in v:
            roles["record_ref"] = letter
    return roles


def _extract_checklist_items(ws, header_row: int, wp_code: str) -> list[dict]:
    items: list[dict] = []
    seq = 0
    for r in range(header_row + 1, ws.max_row + 1):
        a_raw = ws.cell(r, 1).value
        b_raw = ws.cell(r, 2).value
        a = cell_str(a_raw)
        b = cell_str(b_raw)
        text = a or ""
        content = b or a or ""
        if not content:
            continue
        if "签字" in content and ("项目" in content or "合伙人" in content or "负责人" in content):
            break

        is_item = False
        seq_label: str | None = None
        if a and ITEM_RE.match(a):
            is_item = True
            seq_label = ITEM_RE.match(a).group(1)
            content = a
        elif b and a and SEQ_COL_RE.match(str(a_raw).strip()):
            is_item = True
            seq_label = str(a_raw).strip()
            content = b

        if not is_item or len(content) < 8:
            continue
        seq += 1
        items.append({
            "row": r,
            "seq": seq,
            "seq_label": seq_label,
            "item_id": f"{wp_code}-chk-{seq:02d}",
            "content": content,
            "content_preview": content[:160],
            "auto_na_condition": _guess_auto_na(content),
        })
    return items


def _guess_auto_na(content: str) -> str | None:
    if "组成部分注册会计师" in content or "组件" in content:
        return "no_component_auditor"
    if "IT" in content.upper() or "信息技术" in content:
        return "no_it_audit"
    if "内控" in content and "缺陷" in content:
        return None
    return None


def _audit_type_for(wp_code: str) -> str:
    return "internal_control" if wp_code.endswith("-2") else "financial"


def audit_review_workbook(path: Path, wp_code: str, enterprise_variant: str | None) -> dict:
    wb = openpyxl.load_workbook(path, data_only=True, read_only=False)
    checklist_sn, header_row = _find_checklist_sheet(wb)
    ws = wb[checklist_sn]
    column_roles = _detect_column_roles(ws, header_row)
    items = _extract_checklist_items(ws, header_row, wp_code)
    record_sn = _find_record_sheet(wb)
    sign_rows = []
    for r in range(header_row, ws.max_row + 1):
        a = cell_str(ws.cell(r, 1).value) or ""
        if "签字" in a:
            sign_rows.append(r)
    has_cover = "底稿目录" in wb.sheetnames
    wb.close()

    parent = wp_code.split("-")[0]
    meta = ROLE_META.get(parent, {})
    return {
        "wp_code": wp_code,
        "enterprise_variant": enterprise_variant,
        "filename": path.name,
        "runtime": "review-checklist",
        "role": meta.get("role"),
        "role_label": meta.get("role_label"),
        "audit_type": _audit_type_for(wp_code),
        "applicable_categories": meta.get("applicable_categories", []),
        "sheets": {
            "checklist": {
                "name": checklist_sn,
                "header_row": header_row,
                "column_roles": column_roles,
                "item_count": len(items),
                "items": items,
                "sign_rows": sign_rows,
            },
            "record": {"name": record_sn} if record_sn else None,
            "cover": "底稿目录" if has_cover else None,
        },
        "parser_ready": len(items) > 0,
        "notes": f"检查项 {len(items)} 条；item_id {wp_code}-chk-{{seq:02d}}",
    }


def audit_target(wp_code: str, enterprise_variant: str | None) -> dict:
    path = resolve_template_path(wp_code, enterprise_variant)
    entry = audit_review_workbook(path, wp_code, enterprise_variant)
    return entry


def merge_audit(entry: dict) -> None:
    existing: list[dict] = []
    if OUT_AUDIT.exists():
        existing = json.loads(OUT_AUDIT.read_text(encoding="utf-8"))
    key = (entry.get("wp_code"), entry.get("enterprise_variant"))
    existing = [
        e for e in existing
        if (e.get("wp_code"), e.get("enterprise_variant")) != key
    ]
    existing.append(entry)
    existing.sort(key=lambda x: (x.get("wp_code") or "", x.get("enterprise_variant") or ""))
    OUT_AUDIT.write_text(json.dumps(existing, ensure_ascii=False, indent=2), encoding="utf-8")


def write_definitions_from_audit() -> None:
    if not OUT_AUDIT.exists():
        raise FileNotFoundError(f"Run audit first: {OUT_AUDIT}")
    audits = json.loads(OUT_AUDIT.read_text(encoding="utf-8"))
    templates: dict[str, dict] = {}
    auto_na: dict[str, dict] = {
        "no_component_auditor": {"label": "无组成部分注册会计师", "check": "has_component_auditor == false"},
        "no_it_audit": {"label": "无 IT 审计程序", "check": "has_it_audit == false"},
        "not_large_soe": {"label": "非大型国企", "check": "is_large_soe == false"},
        "no_internal_control_audit": {
            "label": "非内控审计",
            "check": "audit_type != 'internal_control' and audit_type != 'combined'",
        },
    }

    for entry in audits:
        if not entry.get("parser_ready"):
            continue
        wp = entry["wp_code"]
        variant = entry.get("enterprise_variant")
        def_key = wp if not variant else f"{wp}:{variant}"
        checklist = entry["sheets"]["checklist"]
        templates[def_key] = {
            "wp_code": wp,
            "enterprise_variant": variant,
            "role": entry.get("role"),
            "role_label": entry.get("role_label"),
            "audit_type": entry.get("audit_type"),
            "audit_type_label": "财务报表审计" if entry.get("audit_type") == "financial" else "内控审计",
            "applicable_categories": entry.get("applicable_categories", []),
            "filename": entry.get("filename"),
            "items": [
                {
                    "seq": it["seq"],
                    "content": it["content"],
                    "auto_na_condition": it.get("auto_na_condition"),
                    "item_id": it["item_id"],
                }
                for it in checklist["items"]
            ],
        }

    payload = {
        "version": "1.0",
        "description": "A21~A25 复核检查项（由 audit_a21_a25_xlsx.py 生成）",
        "source_audit": str(OUT_AUDIT.relative_to(ROOT)),
        "templates": templates,
        "auto_na_conditions": auto_na,
        "role_priority": ["field_lead", "manager", "partner", "quality_reviewer", "eqcr"],
    }
    OUT_DEFS.write_text(json.dumps(payload, ensure_ascii=False, indent=2), encoding="utf-8")
    print(f"definitions -> {OUT_DEFS.relative_to(ROOT)} ({len(templates)} templates)")


def main() -> None:
    parser = argparse.ArgumentParser(description="Audit A21~A25 review xlsx templates")
    parser.add_argument(
        "wp_codes",
        nargs="*",
        help="默认全部 12 条目（含 A24-1/A25-1 双变体）",
    )
    parser.add_argument(
        "--write-definitions",
        action="store_true",
        help="从 audit JSON 生成 a21_a25_review_definitions.json",
    )
    args = parser.parse_args()

    targets = AUDIT_TARGETS
    if args.wp_codes:
        targets = [(c, None) for c in args.wp_codes]

    missing = 0
    for wp_code, variant in targets:
        if args.wp_codes and variant:
            continue
        try:
            entry = audit_target(wp_code, variant)
        except (FileNotFoundError, ValueError) as exc:
            missing += 1
            print(f"[MISSING] {wp_code} ({variant}): {exc}")
            merge_audit({
                "wp_code": wp_code,
                "enterprise_variant": variant,
                "filename": None,
                "parser_ready": False,
                "data_blocked": True,
                "notes": str(exc),
            })
            continue
        merge_audit(entry)
        n = entry["sheets"]["checklist"]["item_count"]
        label = f"{wp_code}" + (f" [{variant}]" if variant else "")
        print(f"audited {label}: {entry['filename']} — {n} items")

    if args.write_definitions:
        write_definitions_from_audit()

    if missing:
        print(f"\n[WARN] {missing} target(s) failed")
    else:
        print(f"\nAudit complete -> {OUT_AUDIT.relative_to(ROOT)}")


if __name__ == "__main__":
    main()

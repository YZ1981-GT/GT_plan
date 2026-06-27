"""Inspect A5-1 template - full detail for all sheets."""
import openpyxl

wb = openpyxl.load_workbook(r"wp_templates/A/A5-1 现金流量表审计.xlsx", read_only=True, data_only=True)
print("Sheets:", wb.sheetnames)
print(f"Total sheets: {len(wb.sheetnames)}")

for sn in wb.sheetnames:
    if sn == 'GT_Custom':
        continue
    ws = wb[sn]
    max_row = ws.max_row or 0
    max_col = ws.max_column or 0
    print(f"\n{'='*70}")
    print(f"Sheet: {sn}")
    print(f"Dimensions: {max_row} rows x {max_col} cols")
    print("="*70)
    # Print ALL rows for complete understanding
    for i, row in enumerate(ws.iter_rows(min_row=1, max_row=max_row, values_only=True), 1):
        cells = [str(c).replace('\xa0', ' ')[:80] if c else "" for c in row[:max_col]]
        line = " | ".join(cells)
        if line.strip(" |"):
            print(f"  R{i:03d}: {line}")
wb.close()

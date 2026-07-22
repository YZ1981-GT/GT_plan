# -*- coding: utf-8 -*-
"""Fix / enhance H2-7 工程造价比较分析表 in H2 Excel template.

改进要点（对齐致同编制思路并补强可操作性）：
1. 审计目标与程序匹配（存在性 + 计价合理性 + 现金流勾稽）
2. 单方造价防 #DIV/0!；新增可比价均值/差异/差异率公式列
3. 增加合计行（单方造价=加权平均）及主体层现金流勾稽汇总
4. 提示补充重大差异阈值与非建筑类工程处理说明
"""
from __future__ import annotations

import shutil
import sys
from pathlib import Path

sys.stdout.reconfigure(encoding="utf-8")

import openpyxl
from openpyxl.comments import Comment
from openpyxl.styles import Alignment, Border, Font, PatternFill, Side
from openpyxl.utils import get_column_letter

SRC = Path(r"d:\GT_plan\backend\wp_templates\H") / "H2 在建工程.xlsx"
TMP = Path(r"d:\GT_plan\backend\wp_templates\H") / "_H2_fixed_tmp.xlsx"

HEADER_FILL = PatternFill("solid", fgColor="D9E2F3")
INPUT_FILL = PatternFill("solid", fgColor="FFF2CC")
FORMULA_FILL = PatternFill("solid", fgColor="E2EFDA")
TOTAL_FILL = PatternFill("solid", fgColor="F2F2F2")
SECTION_FILL = PatternFill("solid", fgColor="FFF2CC")

THIN = Border(
    left=Side(style="thin"),
    right=Side(style="thin"),
    top=Side(style="thin"),
    bottom=Side(style="thin"),
)
BOLD = Font(bold=True)
TITLE_FONT = Font(bold=True, size=14)
WARN = Font(color="C00000", bold=True)


def style_header(cell):
    cell.fill = HEADER_FILL
    cell.font = BOLD
    cell.border = THIN
    cell.alignment = Alignment(horizontal="center", vertical="center", wrap_text=True)


def style_input(cell):
    cell.fill = INPUT_FILL
    cell.border = THIN
    cell.alignment = Alignment(horizontal="right", vertical="center")


def style_formula(cell):
    cell.fill = FORMULA_FILL
    cell.border = THIN
    cell.alignment = Alignment(horizontal="right", vertical="center")


def style_text(cell):
    cell.border = THIN
    cell.alignment = Alignment(horizontal="left", vertical="center", wrap_text=True)


def clear_sheet_body(ws, start_row: int = 5):
    """Clear content from start_row downward within used columns (keep header A1:L4)."""
    # Unmerge ranges that intersect body
    to_unmerge = []
    for mr in list(ws.merged_cells.ranges):
        if mr.min_row >= start_row:
            to_unmerge.append(str(mr))
    for m in to_unmerge:
        ws.unmerge_cells(m)

    max_row = max(ws.max_row, 40)
    max_col = max(ws.max_column, 16)
    for r in range(start_row, max_row + 1):
        for c in range(1, max_col + 1):
            cell = ws.cell(r, c)
            cell.value = None
            cell.fill = PatternFill()
            cell.font = Font()
            cell.border = Border()
            cell.comment = None


def main():
    if not SRC.exists():
        raise FileNotFoundError(SRC)

    shutil.copy2(SRC, TMP)
    wb = openpyxl.load_workbook(TMP)
    ws = wb[[s for s in wb.sheetnames if "H2-7" in s][0]]

    # Expand title merges to cover new columns A:P
    for rng in ("A1:M1", "A2:M2", "A1:P1", "A2:P2"):
        try:
            ws.unmerge_cells(rng)
        except Exception:
            pass
    ws.merge_cells("A1:P1")
    ws.merge_cells("A2:P2")
    ws["A1"] = "致同会计师事务所"
    ws["A2"] = "工程造价比较分析表"
    ws["A1"].font = BOLD
    ws["A2"].font = TITLE_FONT
    ws["A2"].alignment = Alignment(horizontal="center", vertical="center")

    # Keep catalog-linked header cells (A3/D3/I3/L3/M3 etc.)
    ws["L3"] = "索引号："
    ws["M3"] = "H2-7"
    ws["L4"] = "页  次："

    clear_sheet_body(ws, 5)

    # ── 一、审计目标 ──────────────────────────────────────────────────────────
    ws.merge_cells("A5:P5")
    ws["A5"] = (
        "一、审计目标：1.资产负债表中记录的在建工程是存在的，且已记录于恰当的账户；"
        "2.在建工程计价合理，单方造价与市场/同类工程可比价无重大异常差异；"
        "3.本期在建工程现金性投入与现金流量表「购置固定资产、无形资产和其他长期资产支付的现金」"
        "勾稽一致，或差异可合理解释。"
    )
    ws["A5"].alignment = Alignment(wrap_text=True, vertical="center")
    ws.row_dimensions[5].height = 48

    ws["A6"] = "二、审计过程："
    ws["A6"].font = BOLD

    # ── 表头（16列）──────────────────────────────────────────────────────────
    headers = [
        ("A7", "序号"),
        ("B7", "项目"),
        ("C7", "工程项目总造价"),
        ("D7", "建筑面积(㎡)"),
        ("E7", "单方造价"),
        ("F7", "可比价1"),
        ("G7", "可比价2"),
        ("H7", "可比价3"),
        ("I7", "可比价均值"),
        ("J7", "与可比价差异"),
        ("K7", "差异率(%)"),
        ("L7", "可比价来源索引"),
        ("M7", "本期增加额"),
        ("N7", "本期计入现金流量金额"),
        ("O7", "勾稽差异"),
        ("P7", "差异原因/关注事项"),
    ]
    for addr, text in headers:
        ws[addr] = text
        style_header(ws[addr])
    ws.row_dimensions[7].height = 36

    # Sample project rows 8-11 + total row 12
    samples = [
        (8, 1, "工程A"),
        (9, 2, "工程B"),
        (10, 3, "工程C"),
        (11, 4, "……"),
    ]

    for row, seq, name in samples:
        ws.cell(row, 1, seq).border = THIN
        ws.cell(row, 1).alignment = Alignment(horizontal="center")
        ws.cell(row, 2, name)
        style_text(ws.cell(row, 2))

        for col in (3, 4, 6, 7, 8, 12, 13, 14, 16):  # inputs (C D F G H L M N P)
            style_input(ws.cell(row, col))
            if col in (3, 4, 6, 7, 8, 13, 14):
                ws.cell(row, col).number_format = "#,##0.00"

        # E 单方造价
        ws.cell(row, 5, f'=IF(OR(D{row}="",D{row}=0),"",C{row}/D{row})')
        style_formula(ws.cell(row, 5))
        ws.cell(row, 5).number_format = "#,##0.00"
        ws.cell(row, 5).comment = Comment("单方造价=总造价/建筑面积；面积为空时不显示，避免#DIV/0!", "H2-7")

        # I 可比价均值
        ws.cell(row, 9, f'=IF(COUNT(F{row}:H{row})=0,"",AVERAGE(F{row}:H{row}))')
        style_formula(ws.cell(row, 9))
        ws.cell(row, 9).number_format = "#,##0.00"

        # J 与可比价差异
        ws.cell(row, 10, f'=IF(OR(E{row}="",I{row}=""),"",E{row}-I{row})')
        style_formula(ws.cell(row, 10))
        ws.cell(row, 10).number_format = "#,##0.00"

        # K 差异率
        ws.cell(row, 11, f'=IF(OR(I{row}="",I{row}=0),"",J{row}/I{row})')
        style_formula(ws.cell(row, 11))
        ws.cell(row, 11).number_format = "0.00%"
        ws.cell(row, 11).comment = Comment("差异率=与可比价差异/可比价均值；|差异率|>15%建议说明并评估是否需追加程序", "H2-7")

        # O 勾稽差异
        ws.cell(row, 15, f'=IF(AND(M{row}="",N{row}=""),"",N(M{row})-N(N{row}))')
        style_formula(ws.cell(row, 15))
        ws.cell(row, 15).number_format = "#,##0.00"
        ws.cell(row, 15).comment = Comment(
            "勾稽差异=本期增加额-本期计入现金流量金额；常见解释：应付工程款/预付/非现金投入/利息资本化等",
            "H2-7",
        )

        style_text(ws.cell(row, 12))
        style_text(ws.cell(row, 16))

    # 合计行
    total_row = 12
    ws.cell(total_row, 1, "").border = THIN
    ws.cell(total_row, 2, "合计")
    ws.cell(total_row, 2).font = BOLD
    ws.cell(total_row, 2).fill = TOTAL_FILL
    ws.cell(total_row, 2).border = THIN

    for col, letter in [(3, "C"), (4, "D"), (13, "M"), (14, "N"), (15, "O")]:
        cell = ws.cell(total_row, col, f"=SUM({letter}8:{letter}11)")
        cell.fill = TOTAL_FILL
        cell.font = BOLD
        cell.border = THIN
        cell.number_format = "#,##0.00"

    # 合计单方造价 = 加权平均
    ws.cell(total_row, 5, f'=IF(OR(D{total_row}="",D{total_row}=0),"",C{total_row}/D{total_row})')
    style_formula(ws.cell(total_row, 5))
    ws.cell(total_row, 5).fill = TOTAL_FILL
    ws.cell(total_row, 5).font = BOLD
    ws.cell(total_row, 5).number_format = "#,##0.00"
    ws.cell(total_row, 5).comment = Comment("合计行单方造价按总造价/总面积加权，勿对各行单方造价简单平均", "H2-7")

    for col in (6, 7, 8, 9, 10, 11, 12, 16):
        cell = ws.cell(total_row, col, "")
        cell.fill = TOTAL_FILL
        cell.border = THIN

    # ── （二）主体层现金流勾稽汇总 ────────────────────────────────────────────
    ws.merge_cells("A14:P14")
    ws["A14"] = "（二）与现金流量表勾稽汇总（主体层面，通常比逐项分摊更可靠）"
    ws["A14"].font = BOLD
    ws["A14"].fill = SECTION_FILL

    recon_headers = [
        (15, 1, "索引"),
        (15, 2, "项目"),
        (15, 3, "金额"),
        (15, 4, "来源/说明"),
    ]
    for r, c, t in recon_headers:
        cell = ws.cell(r, c, t)
        style_header(cell)

    ws["A16"] = "①"
    ws["B16"] = "本期在建工程增加额合计（上表M列）"
    ws["C16"] = f"=M{total_row}"
    style_formula(ws["C16"])
    ws["C16"].number_format = "#,##0.00"
    ws["D16"] = "来自本表合计；可与H2-2增加合计交叉核对"
    style_text(ws["D16"])

    ws["A17"] = "②"
    ws["B17"] = "现金流量表：购置固定资产、无形资产和其他长期资产支付的现金"
    style_input(ws["C17"])
    ws["C17"].number_format = "#,##0.00"
    ws["C17"].comment = Comment("手工填入现金流量表对应项目金额（注意：通常含购建固定资产等，不限于在建工程）", "H2-7")
    ws["D17"] = "取自现金流量表；范围通常宽于在建工程，差异需分层解释"
    style_text(ws["D17"])

    ws["A18"] = "③"
    ws["B18"] = "勾稽差异（①−②）"
    ws["C18"] = '=IF(C17="","" ,C16-C17)'
    style_formula(ws["C18"])
    ws["C18"].number_format = "#,##0.00"
    ws["C18"].font = WARN
    style_input(ws["D18"])
    ws["D18"] = ""
    ws["D18"].comment = Comment(
        "常见差异原因：应付/预付工程款变动、购入已达可使用状态固定资产、无形资产购置、"
        "非现金投入、利息资本化（非付现）、合并范围/口径差异等",
        "H2-7",
    )

    # ── 三、四 审计说明/结论 ──────────────────────────────────────────────────
    ws.merge_cells("A20:P20")
    ws["A20"] = "三、审计说明："
    ws["A20"].font = BOLD
    ws.merge_cells("A21:P22")
    ws["A21"] = ""
    ws["A21"].alignment = Alignment(wrap_text=True, vertical="top")
    ws["A21"].fill = INPUT_FILL
    ws.row_dimensions[21].height = 40
    ws.row_dimensions[22].height = 40

    ws.merge_cells("A23:P23")
    ws["A23"] = "四、审计结论："
    ws["A23"].font = BOLD
    ws.merge_cells("A24:P25")
    ws["A24"] = ""
    ws["A24"].alignment = Alignment(wrap_text=True, vertical="top")
    ws["A24"].fill = INPUT_FILL
    ws.row_dimensions[24].height = 36
    ws.row_dimensions[25].height = 36

    # ── 提示 ─────────────────────────────────────────────────────────────────
    ws.merge_cells("A27:P27")
    ws["A27"] = "提示："
    ws["A27"].font = BOLD

    tips = [
        (28, "1.通过募投项目可行性研究报告、工程概预算/合同等，获取工程项目总造价及建筑面积；"
             "计算单方造价，并与不少于一项可比价（同类工程/造价信息价/评估报告）比较；"
             "填列可比价来源索引以便复核。"),
        (29, "2.单方造价与可比价均值差异率绝对值＞15%（或超过项目组确定的分析阈值）时，"
             "应在「差异原因/关注事项」说明，并评估是否需补充程序（如询价、复核变更签证、关注关联方定价）。"),
        (30, "3.无建筑面积的装置类/管网类工程：建筑面积可留空，改在备注说明比较口径"
             "（如单位产能造价、单位长度造价），勿强行填0导致公式失真。"),
        (31, "4.对比本期在建工程增加（现金性投入）与现金流量表「购置固定资产、无形资产和其他长期资产支付的现金」；"
             "优先做主体层汇总勾稽（见上表②③），项目层分摊仅在可可靠归集时使用。"),
        (32, "5.勾稽差异常见解释：应付工程款/预付账款变动、购入不经在建工程的固定资产、"
             "无形资产购置、非现金投入、资本化利息（非付现）等；重大非预期差异应与客户讨论。"),
        (33, "6.确定上述结果是否表明需要补充审计程序；如是，执行适当应对程序，并在审计说明中记录。"),
    ]
    for row, text in tips:
        ws.merge_cells(f"A{row}:P{row}")
        ws[f"A{row}"] = text
        ws[f"A{row}"].alignment = Alignment(wrap_text=True, vertical="center")
        ws.row_dimensions[row].height = 32

    # Column widths
    widths = {
        "A": 6, "B": 18, "C": 14, "D": 12, "E": 11,
        "F": 10, "G": 10, "H": 10, "I": 11, "J": 12,
        "K": 10, "L": 14, "M": 12, "N": 16, "O": 11, "P": 18,
    }
    for col, w in widths.items():
        ws.column_dimensions[col].width = w

    alt = SRC.with_name(SRC.stem + "_H2-7updated.xlsx")
    wb.save(TMP)
    try:
        shutil.copy2(TMP, SRC)
        print(f"OK: updated {SRC}")
    except PermissionError:
        shutil.copy2(TMP, alt)
        print("源文件被 Excel/WPS 占用，无法覆盖。")
        print(f"已另存为: {alt}")
        print("请关闭占用后执行: python backend/scripts/_fix_h2_7_template.py")
        print(f"或手动用该文件替换: {SRC.name}")


if __name__ == "__main__":
    main()

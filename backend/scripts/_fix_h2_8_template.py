# -*- coding: utf-8 -*-
"""改进 H2-8 在建工程增加检查表模板。

主要改动：
1. 审计目标聚焦「本期增加」细节测试认定
2. 样本选取补充计划覆盖率、与 H2-2 勾稽说明
3. 测试内容补全资本化/勾稽/关联方要点；U列由「……」改为「资本化判断」
4. 检查比例公式防 DIV/0；审计说明/结论给出可执行指引
5. 附注去重：使用说明 + 执行核对清单（含交叉索引）+ 舞弊提示
"""
from __future__ import annotations

import sys
from pathlib import Path

sys.stdout.reconfigure(encoding="utf-8")

from openpyxl import load_workbook
from openpyxl.comments import Comment
from openpyxl.styles import Alignment, Font, PatternFill

SRC = Path(r"d:\GT_plan\backend\wp_templates\H") / "H2 在建工程.xlsx"
TMP = Path(r"d:\GT_plan\backend\wp_templates\H") / "_H2_8_fixed_tmp.xlsx"

BLUE = Font(name="Arial Narrow", size=10, color="0000FF")
RED = Font(name="Arial Narrow", size=10, color="FF0000")
BLACK = Font(name="Arial Narrow", size=10)
BLACK_BOLD = Font(name="Arial Narrow", size=10, bold=True)
SONG = Font(name="宋体", size=10)
SONG_BOLD = Font(name="宋体", size=10, bold=True)
HEADER_BOLD = Font(name="宋体", size=10, bold=True)
WRAP = Alignment(wrap_text=True, vertical="center")
WRAP_TOP = Alignment(wrap_text=True, vertical="top")


def main() -> None:
    wb = load_workbook(SRC)
    ws = wb[[s for s in wb.sheetnames if "H2-8" in s or "增加检查" in s][0]]

    # ── 一、审计目标：聚焦本期增加细节测试 ─────────────────────────────
    ws["A5"] = (
        "一、审计目标：针对本期在建工程增加实施细节测试——"
        "（1）存在/发生：抽查的增加真实发生，对应工程项目存在；"
        "（2）准确性与计价：金额与合同/进度/发票等证据相符，资本化范围恰当（应资本化 vs 应费用化）；"
        "（3）截止：记入正确的会计期间；"
        "（4）权利与义务：相关合同权利义务归属于被审计单位。"
        "重大遗漏风险结合总体分析程序与 H2-2/H2-6 关注完整性。"
    )
    ws["A5"].font = BLACK_BOLD
    ws["A5"].alignment = WRAP
    ws.row_dimensions[5].height = 48

    # ── 二、样本选取：补计划覆盖率、与明细表勾稽 ───────────────────────
    ws["A6"] = "二、样本选取标准与规模："
    ws["A6"].font = BLACK_BOLD

    ws["A7"] = "测试总体："
    ws["B7"] = "本期在建工程借方发生额（与 H2-2 本期增加合计勾稽），共 XX 笔、金额 XX"
    ws["B7"].font = BLUE
    ws["F7"] = "特定样本："
    ws["G7"] = (
        "超过重要性/明显重大金额、关联方交易、异常对方科目/摘要、年末集中入账等全部详查，共 XX 笔、金额 XX"
    )
    ws["G7"].font = BLUE

    ws["A8"] = "抽样总体："
    ws["B8"] = "测试总体扣除特定样本后剩余项目，共 XX 笔、金额 XX"
    ws["B8"].font = BLUE
    ws["F8"] = "确定的抽样样本量："
    ws["G8"] = "抽取 XX 笔、金额 XX；计划检查比例 XX%（样本金额/本期增加合计）"
    ws["G8"].font = BLUE
    ws["H8"] = "（若使用样本计算器，计算过程见 <XX> 底稿）"
    ws["H8"].font = BLUE

    ws["A9"] = "抽样方法："
    ws["B9"] = "随机选样 / 系统选样 / 货币单元抽样 / 随意选样（非统计抽样）"
    ws["B9"].font = BLUE
    ws["F9"] = "抽样过程："
    ws["G9"] = (
        "使用 IDEA（或其他抽样工具）选取样本；抽样过程与结果见 <XX> 底稿。"
        "特定样本与抽样样本合计构成检查总体。"
    )
    ws["G9"].font = BLUE

    # ── 三、测试内容说明 ───────────────────────────────────────────────
    ws["A10"] = "三、测试："
    ws["A10"].font = BLACK_BOLD
    ws["A11"] = (
        "测试内容说明：1.原始凭证是否齐全；"
        "2.记账凭证与原始凭证是否相符；"
        "3.账务处理是否正确（含资本化/费用化划分、对方科目）；"
        "4.是否记录于恰当会计期间；"
        "5.合同金额、形象进度/监理报告、结算及付款是否勾稽一致；"
        "6.大额或关联交易是否单独识别（→H2-17）。"
        "不适用的证据列填 N/A；证据列可按被审计单位情况增删。"
    )
    ws["A11"].font = RED
    ws["A11"].alignment = WRAP
    ws.row_dimensions[11].height = 36

    # ── 表头：U列由「……」改为「资本化判断」；W列明确为检查结论 ────────
    ws["U12"] = "资本化判断"
    ws["U12"].font = HEADER_BOLD
    ws["U12"].alignment = Alignment(horizontal="center", vertical="center", wrap_text=True)
    ws["U12"].comment = Comment(
        "Y=应计入在建工程成本且本期资本化恰当；N=应费用化或其他科目；"
        "N/A=本样本不适用（如纯物资调拨等）。判断依据写入审计说明或备注。",
        "H2-8",
        width=280,
        height=80,
    )

    ws["W12"] = "检查结论"
    ws["W12"].font = HEADER_BOLD
    ws["W12"].alignment = Alignment(horizontal="center", vertical="center", wrap_text=True)
    ws["W12"].comment = Comment(
        "填：无异常 / 存疑 / 需调整。需调整事项交叉索引至 H2-3。",
        "H2-8",
        width=220,
        height=50,
    )

    # 右侧提示（原 X13）
    ws["X13"] = "检查的关键证据和要素根据被审计单位具体情况修改；不适用列填 N/A"
    ws["X13"].font = BLUE
    ws["X13"].alignment = WRAP

    # 监理/领料/发票表头保留红色提示色（提醒可按业态裁剪）
    for coord in ("L12", "N12", "R12", "L13", "M13", "N13", "O13", "P13", "Q13", "R13", "S13", "T13"):
        if ws[coord].value:
            ws[coord].font = Font(name="宋体", size=10, bold=True, color="FF0000")

    # ── 汇总行：防 DIV/0；补充说明 ─────────────────────────────────────
    ws["A30"] = "合计（已检查样本）"
    ws["A31"] = "本期增加在建工程合计（H2-2）"
    ws["H31"].comment = Comment(
        "取自明细表 H2-2 本期增加合计，用于计算检查比例。",
        "H2-8",
        width=200,
        height=40,
    )
    ws["A32"] = "检查比例"
    ws["H32"] = '=IF(OR(H31="",H31=0),"",H30/H31)'
    ws["H32"].number_format = "0.00%"
    ws["H32"].comment = Comment(
        "已检查样本金额合计 ÷ 本期增加合计。低于计划比例须扩大样本或在审计说明中解释。",
        "H2-8",
        width=260,
        height=55,
    )

    # ── 四、审计说明指引 ───────────────────────────────────────────────
    ws["A33"] = "四、审计说明："
    ws["A33"].font = BLACK_BOLD
    ws["A34"] = (
        "编制指引：①样本构成（特定样本+抽样样本）、抽样方法及计划/实际检查比例；"
        "比例偏低须扩大样本或说明原因。"
        "②逐项异常/存疑/需调整事项及证据索引；调整分录→H2-3。"
        "③利息资本化详见 H2-10/H2-11；关联方交易详见 H2-17。"
    )
    ws["A34"].font = RED
    ws["A34"].alignment = WRAP
    ws.row_dimensions[34].height = 36
    ws["A35"] = ""
    ws["A36"] = None
    ws["A37"] = None

    # ── 五、审计结论：给出可填写模板 ───────────────────────────────────
    ws["A38"] = "五、审计结论："
    ws["A38"].font = BLACK_BOLD
    ws["A39"] = (
        "基于上述检查，本期抽查的在建工程增加在存在性、准确性/资本化划分及截止方面"
        "未见重大异常 / 发现以下需调整事项（详见审计说明及 H2-3）：________。"
        "检查比例 ___%，可为相关认定提供充分、适当的审计证据。"
    )
    ws["A39"].font = BLUE
    ws["A39"].alignment = WRAP
    ws.row_dimensions[39].height = 36

    # ── 附注：去重并改为「使用说明 + 核对清单 + 舞弊提示」──────────────
    # 清除旧 A40–A57 文本后重写（保留行数，避免破坏过多版式）
    for r in range(40, 58):
        ws.cell(row=r, column=1).value = None
        ws.cell(row=r, column=1).font = SONG
        ws.cell(row=r, column=1).alignment = WRAP_TOP

    ws["A40"] = "附注"
    ws["A40"].font = SONG_BOLD

    ws["A41"] = "【使用说明】"
    ws["A41"].font = BLACK_BOLD
    ws["A42"] = (
        "1.本表汇总本期在建工程增加的细节测试。测试总体金额应与 H2-2 本期增加合计、H2-6 借方发生额勾稽。"
    )
    ws["A43"] = (
        "2.取证范围：立项/批复、施工（采购）合同、进度确认/监理报告、工程结算单、发票、领料出库、付款回单及授权审批等；"
        "按增加方式选用证据列，不适用填 N/A。"
    )
    ws["A44"] = (
        "3.出包工程重进度与结算勾稽；自营工程重领料与人工/机械归集；设备购置重合同、发票与验收；"
        "利息/待摊支出分别交叉 H2-10/H2-11 及费用类科目，避免在本表重复详测。"
    )

    ws["A45"] = "【执行核对清单】（完成情况在「审计说明」中简述，打勾仅为编制辅助）"
    ws["A45"].font = BLACK_BOLD
    ws["A46"] = "□ 工程款支付与合同、进度/监理报告及授权审批一致，会计处理正确"
    ws["A47"] = "□ 工程物资领用有审批，会计处理正确"
    ws["A48"] = "□ 借款费用资本化条件、方法、期间、金额正确（交叉 H2-10、H2-11）"
    ws["A49"] = "□ 工程管理费等资本化金额合理（交叉管理费用等科目）"
    ws["A50"] = "□ 土地使用权与地上建筑物分别核算（土地→无形资产，不挤占在建工程）"
    ws["A51"] = "□ 出包工程按合理估计进度及合同结算款入账；设备交付建造承包商时点恰当"
    ws["A52"] = "□ 工程物资盘盈盘亏、报废毁损净损益及残料/废料处置冲减成本处理正确"
    ws["A53"] = (
        "□ 待摊支出（管理费、征地、可研、临时设施、监理、税费等）分配合理；"
        "达预定可使用状态时转入对应成本；付现与现金流量表「购建固定资产、无形资产和其他长期资产支付的现金」无重大未解释差异"
    )
    ws.row_dimensions[53].height = 32

    ws["A54"] = "【舞弊与关联方提示】"
    ws["A54"].font = BLACK_BOLD
    ws["A55"] = (
        "参照《中国注册会计师审计准则问题解答第18号——识别和应对第三方配合实施财务舞弊》："
        "存在大额长期资产建造/采购时，关注是否通过虚增在建工程造价套取资金；"
        "若同时存在关联方资金占用或其他套取资金舞弊风险因素，应评价是否存在第三方配合舞弊并实施追加程序。"
        "已识别关联方增加应标记并交叉至 H2-17。"
    )
    ws["A55"].font = BLUE
    ws["A55"].alignment = WRAP
    ws.row_dimensions[55].height = 48
    ws["A56"] = None
    ws["A57"] = None

    # 略增关键说明行高
    for r in (42, 43, 44):
        ws.row_dimensions[r].height = 28

    try:
        wb.save(SRC)
        print(f"OK saved: {SRC}")
    except PermissionError:
        wb.save(TMP)
        print(f"源文件被占用，已保存临时文件: {TMP}")
        print("请关闭 Excel/WPS 中的「H2 在建工程.xlsx」后，将临时文件覆盖回原文件。")
        raise SystemExit(2)


if __name__ == "__main__":
    main()

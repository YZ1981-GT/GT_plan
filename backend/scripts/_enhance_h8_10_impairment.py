# -*- coding: utf-8 -*-
"""Enhance H8-10 使用权资产减值测算表 Excel template.

源模板问题：
1. 表头公式标注 H9「⑥＝⑤－②」与单元格公式 IF(D>G,D-G,0) 矛盾（且会算出负数）
2. ⑧=⑥-⑦ 允许为负，与「减值一经确认不得转回」(CAS8) 冲突
3. 账面价值②未注明是否含已计提减值，易与审定表净值混淆
4. 无减值迹象时仍可填可收回金额并出减值，未按 CAS8「先迹象、后测试」闸门
5. 审计目标偏笼统；提示中「转销」易与「转回」混淆；未索引 H8-11/H8-8(含减值)
6. 审计结论区空白，无结构化结论模板；E/F 与 H8-11 无勾稽说明

改进：
- 统一⑥/⑧公式文案与单元格公式（MAX 闸门）
- B=无迹象时强制⑤⑥⑧=0；多提金额改记「多提待查(不得转回)」列
- 明确②口径；强化目标/提示/结论；E/F 备注链 H8-11
"""
from __future__ import annotations

from pathlib import Path

import openpyxl
from openpyxl.comments import Comment
from openpyxl.styles import Alignment, Border, Font, PatternFill, Side
from openpyxl.utils import get_column_letter
from openpyxl.worksheet.datavalidation import DataValidation

ROOT = Path(r"D:/GT_plan")
TARGETS = [ROOT / "backend" / "wp_templates" / "H" / "H8 使用权资产.xlsx"]
base = ROOT / "基础数据"
if base.exists():
    for p in base.rglob("H8 使用权资产.xlsx"):
        if p not in TARGETS:
            TARGETS.append(p)

SHEET = "减值测算表H8-10"
DATA_START, DATA_END = 10, 21
TOTAL_ROW = 22

THIN = Border(
    left=Side(style="thin", color="B0B0B0"),
    right=Side(style="thin", color="B0B0B0"),
    top=Side(style="thin", color="B0B0B0"),
    bottom=Side(style="thin", color="B0B0B0"),
)
HEADER_FILL = PatternFill("solid", fgColor="DDEBF7")
FORMULA_FILL = PatternFill("solid", fgColor="E2EFDA")
TOTAL_FILL = PatternFill("solid", fgColor="FFF2CC")
WARN_FILL = PatternFill("solid", fgColor="FCE4D6")
NOTE_FILL = PatternFill("solid", fgColor="FFF8E7")
GUIDE_FONT = Font(color="0563C1", size=9)
BOLD = Font(bold=True, size=10)
HEADER_FONT = Font(bold=True, size=9)
OBJ_FONT = Font(size=10)
TITLE_FONT = Font(bold=True, size=14, color="1F4E79")
WRAP = Alignment(wrap_text=True, vertical="center", horizontal="center")
LEFT_WRAP = Alignment(wrap_text=True, vertical="center", horizontal="left")


def _unmerge_overlapping(ws, min_row: int, max_row: int, min_col: int, max_col: int) -> None:
    to_remove = []
    for mr in list(ws.merged_cells.ranges):
        if not (mr.max_row < min_row or mr.min_row > max_row or mr.max_col < min_col or mr.min_col > max_col):
            to_remove.append(str(mr))
    for s in to_remove:
        ws.unmerge_cells(s)


def enhance(path: Path, save_as: Path | None = None) -> None:
    wb = openpyxl.load_workbook(path)
    if SHEET not in wb.sheetnames:
        raise SystemExit(f"sheet not found: {SHEET} in {path}")
    ws = wb[SHEET]
    out = save_as or path

    # ── 1. 审计目标（聚焦减值认定）──────────────────────────────────────────
    ws["A5"] = (
        "一、审计目标：确认使用权资产减值准备以恰当金额计入财务报表；"
        "减值迹象识别充分，可收回金额=max(公允减处置费用,预计未来现金流量现值)测算合理；"
        "本期补提/结转会计处理正确；减值一经确认不得转回（CAS8）；相关披露恰当。"
    )
    ws["A5"].font = OBJ_FONT
    ws["A5"].alignment = LEFT_WRAP

    # ── 2. 表头：明确②口径 + 修正⑥文案 + 新增「多提待查」列 M ──────────────
    _unmerge_overlapping(ws, 1, 2, 1, 13)
    _unmerge_overlapping(ws, 7, 9, 1, 13)
    try:
        ws.merge_cells("A1:M1")
        ws.merge_cells("A2:M2")
    except ValueError:
        pass

    # 重新合并表头（A–L 原结构 + M）
    merges = [
        "A7:A9",
        "B7:B9",
        "C7:C8",
        "D7:D8",
        "E7:G7",
        "H7:H8",
        "I7:I8",
        "J7:J8",
        "K7:K9",
        "L7:L9",
        "M7:M8",
    ]
    for m in merges:
        try:
            ws.merge_cells(m)
        except ValueError:
            pass

    headers = {
        "A7": "使用权资产名称",
        "B7": "是否存在减值迹象",
        "C7": "减值迹象描述",
        "D7": "账面价值\n（原值−累计折旧，不含减值）",
        "E7": "可收回金额",
        "H7": "期末应计提的减值金额\n（当⑤＜②）",
        "I7": "期末账面已计提的减值准备",
        "J7": "本期应补提的减值准备\n（减值一经确认不得转回）",
        "K7": "工作底稿索引号",
        "L7": "备注",
        "M7": "多提待查金额\n（⑦＞⑥，不得转回）",
        "E8": "公允价值减处置费用的净额",
        "F8": "预计未来现金流量现值",
        "G8": "可收回金额（③、④较高者）",
        "C9": "①",
        "D9": "②",
        "E9": "③",
        "F9": "④",
        "G9": "⑤=MAX(③,④)",
        "H9": "⑥=MAX(②−⑤,0)",
        "I9": "⑦",
        "J9": "⑧=MAX(⑥−⑦,0)",
        "M9": "⑨=MAX(⑦−⑥,0)",
    }
    for coord, text in headers.items():
        cell = ws[coord]
        cell.value = text
        cell.font = HEADER_FONT
        cell.fill = HEADER_FILL
        cell.alignment = WRAP
        cell.border = THIN

    # 子列表头边框
    for col in range(1, 14):
        for r in (7, 8, 9):
            c = ws.cell(r, col)
            c.border = THIN
            if c.value is None and r == 8 and col in (1, 2, 3, 4, 8, 9, 10, 11, 12):
                c.fill = HEADER_FILL

    # 列宽
    widths = {
        "A": 18, "B": 12, "C": 18, "D": 14, "E": 12, "F": 12, "G": 12,
        "H": 14, "I": 14, "J": 14, "K": 12, "L": 14, "M": 12,
    }
    for col, w in widths.items():
        ws.column_dimensions[col].width = w
    ws.row_dimensions[7].height = 36
    ws.row_dimensions[8].height = 30

    # 批注：口径与跨表
    ws["D7"].comment = Comment(
        "口径：原值−累计折旧，不含减值准备。可与明细表H8-2/审定表H8-1原值、累计折旧勾稽；"
        "勿直接填审定净值（净值已扣减值）。",
        "H8-10",
    )
    ws["E8"].comment = Comment(
        "取自可收回金额测试表H8-11「公允价值−处置费用」；无活跃市场时可留空，以④为准。",
        "H8-10",
    )
    ws["F8"].comment = Comment(
        "取自H8-11 DCF现值；无减值迹象（B=×）时勿填，公式将强制⑤⑥⑧=0。",
        "H8-10",
    )
    ws["J7"].comment = Comment(
        "CAS8：资产减值损失一经确认，在以后会计期间不得转回。"
        "若⑦＞⑥，差额记入⑨「多提待查」，查明是否处置结转遗漏等，禁止做转回分录。",
        "H8-10",
    )

    # ── 3. 数据行公式（迹象闸门 + MAX 不得转回）────────────────────────────
    # 清除旧 DV 后重建
    ws.data_validations.dataValidation = [
        dv for dv in ws.data_validations.dataValidation
        if "B10" not in str(dv.sqref) and "B10:B21" not in str(dv.sqref)
    ]
    dv_sign = DataValidation(
        type="list",
        formula1='"√,×"',
        allow_blank=True,
        showErrorMessage=True,
        errorTitle="减值迹象",
        error="请选择 √（存在）或 ×（不存在）",
    )
    dv_sign.add(f"B{DATA_START}:B{DATA_END}")
    ws.add_data_validation(dv_sign)

    for r in range(DATA_START, DATA_END + 1):
        # ⑤ 可收回金额：无迹象→0；有迹象→MAX(③,④)
        ws[f"G{r}"] = f'=IF(OR(B{r}="",B{r}="×"),0,MAX(E{r},F{r}))'
        # ⑥ 应计提：无迹象→0；否则 MAX(②−⑤,0)
        ws[f"H{r}"] = f'=IF(OR(B{r}="",B{r}="×"),0,MAX(D{r}-G{r},0))'
        # ⑧ 本期补提：不得为负
        ws[f"J{r}"] = f"=MAX(H{r}-I{r},0)"
        # ⑨ 多提待查（不得转回）
        ws[f"M{r}"] = f"=MAX(I{r}-H{r},0)"

        for col in ("G", "H", "J", "M"):
            cell = ws[f"{col}{r}"]
            cell.fill = FORMULA_FILL
            cell.border = THIN
            cell.number_format = "#,##0.00"
        for col in ("A", "B", "C", "D", "E", "F", "I", "K", "L"):
            ws[f"{col}{r}"].border = THIN
            if col in ("D", "E", "F", "I"):
                ws[f"{col}{r}"].number_format = "#,##0.00"

    # 合计行
    ws[f"A{TOTAL_ROW}"] = "合计"
    ws[f"A{TOTAL_ROW}"].font = BOLD
    ws[f"A{TOTAL_ROW}"].fill = TOTAL_FILL
    for col in ("D", "E", "F", "G", "H", "I", "J", "M"):
        cell = ws[f"{col}{TOTAL_ROW}"]
        cell.value = f"=SUM({col}{DATA_START}:{col}{DATA_END})"
        cell.fill = TOTAL_FILL
        cell.border = THIN
        cell.font = BOLD
        cell.number_format = "#,##0.00"
    for col in ("B", "C", "K", "L"):
        ws[f"{col}{TOTAL_ROW}"].fill = TOTAL_FILL
        ws[f"{col}{TOTAL_ROW}"].border = THIN

    # 多提合计警示底色
    ws[f"M{TOTAL_ROW}"].fill = WARN_FILL

    # ── 4. 审计说明 / 结论模板 ──────────────────────────────────────────────
    ws["A23"] = "三、审计说明："
    ws["A23"].font = BOLD
    ws["A24"] = (
        "编制说明：①先评减值迹象（B列）；有迹象者完成H8-11可收回金额测试后回填③④；"
        "②账面价值取自H8-2原值−累计折旧（不含减值）；⑦取自账面减值准备科目余额；"
        "⑧>0需建议补提并考虑切换「折旧测算表（含减值）H8-8」；⑨>0须查多提原因（处置结转等），禁止转回。"
    )
    ws["A24"].font = GUIDE_FONT
    ws["A24"].alignment = LEFT_WRAP
    ws["A24"].fill = NOTE_FILL
    ws.row_dimensions[24].height = 48

    ws["A26"] = "四、审计结论："
    ws["A26"].font = BOLD
    ws["A27"] = (
        "经审计：（1）减值迹象识别□充分 / □需补充；（2）可收回金额测算□合理 / □需调整；"
        "（3）本期应补提合计（⑧）______元，已建议调整□是 / □否 / □不适用；"
        "（4）多提待查（⑨）______元，原因及处理____________________；"
        "（5）使用权资产减值准备在重大方面□公允反映 / □存在错报风险。"
    )
    ws["A27"].font = OBJ_FONT
    ws["A27"].alignment = LEFT_WRAP
    ws.row_dimensions[27].height = 48

    # ── 5. 提示（纠正转销/转回表述 + 跨表闭环）──────────────────────────────
    ws["A30"] = "提示："
    ws["A30"].font = BOLD
    tips = [
        (
            "A31",
            "1.获取或编制使用权资产减值准备明细表，复核加计正确，并与总账、明细账及审定表H8-1减值准备余额核对相符；"
            "账面价值②与明细表H8-2原值、累计折旧勾稽。",
        ),
        (
            "A32",
            "2.检查减值准备计提的批准程序及书面依据是否充分；会计处理是否正确。"
            "注意：CAS8规定减值损失一经确认不得转回；资产处置/转让时相应减值准备应一并结转（终止确认），"
            "结转≠转回，勿混淆。",
        ),
        (
            "A33",
            "3.存在减值迹象的项目，须完成可收回金额测试表H8-11（公允净额及/或DCF），将③④回填本表；"
            "无减值迹象（B=×）不得仅凭估计计提减值。",
        ),
        (
            "A34",
            "4.重新测试管理层减值估计的合理性，比较预期与已记录金额，调查重大差异；"
            "本期补提后应切换至「折旧测算表（含减值）H8-8」重算后续折旧，并与H8-9分配分析勾稽。",
        ),
        (
            "A35",
            "5.检查上期已确认减值准备本期变动是否适当（处置结转、重分类等）；"
            "若⑨多提待查>0，查明原因并评估是否需调整，禁止以「转回」冲减资产减值损失。",
        ),
    ]
    for coord, text in tips:
        ws[coord] = text
        ws[coord].font = GUIDE_FONT
        ws[coord].alignment = LEFT_WRAP
        ws.row_dimensions[int(coord[1:])].height = 36

    # 标题样式
    ws["A2"].font = TITLE_FONT

    try:
        wb.save(out)
        print(f"OK: {out}")
    except PermissionError:
        alt = out.with_name(out.stem + "_H810enhanced.xlsx")
        wb.save(alt)
        print(f"LOCKED {out} → saved {alt}")


def main() -> None:
    for p in TARGETS:
        if not p.exists():
            print(f"SKIP missing: {p}")
            continue
        # 跳过已是增强副本，避免循环套娃
        if "_H810enhanced" in p.name or "_H813enhanced" in p.name:
            continue
        enhance(p)
    # 若主模板被锁，对已生成的增强副本再跑一遍标题合并等补丁
    enh = ROOT / "backend" / "wp_templates" / "H" / "H8 使用权资产_H810enhanced.xlsx"
    if enh.exists():
        enhance(enh, save_as=enh)


if __name__ == "__main__":
    main()

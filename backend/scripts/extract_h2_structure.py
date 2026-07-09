"""
Phase 0 Task 0.1: openpyxl脚本读取H2在建工程.xlsx全部21 sheet
提取：sheet名/列头/行数/公式单元格/合并区域/数据类型
产出：h2_structure_summary.json（权威列头+公式清单）
验证：21 sheet结构与本spec描述一致
"""

import json
import os
from pathlib import Path

import openpyxl
from openpyxl.utils import get_column_letter

# H2在建工程.xlsx路径
XLSX_PATH = Path(
    r"d:\GT_plan\基础数据\致同通用审计程序及底稿模板（2025年修订）"
    r"\1.致同审计程序及底稿模板（2025年）"
    r"\4.风险应对-实质性程序（D-N）"
    r"\H 固定资产循环"
    r"\H2 在建工程.xlsx"
)

OUTPUT_PATH = Path(r"d:\GT_plan\.kiro\specs\h2-construction-in-progress\h2_structure_summary.json")

# spec中定义的21个sheet（按预期）
EXPECTED_SHEETS = [
    "Tab_Index",
    "Procedure_Table_H2A",
    "Adjudication_H2_1",
    "Disclosure_Listed",
    "Disclosure_SOE",
    "Detail_H2_2",
    "Adjustment_H2_3",
    "Analysis_H2_4",
    "Transfer_Check_H2_5",
    "Review_Record_H2_6",
    "Cost_Comparison_H2_7",
    "Addition_Check_H2_8",
    "Decrease_Check_H2_9",
    "Interest_Cap_NoBorrow_H2_10",
    "Interest_Cap_WithBorrow_H2_11",
    "Stocktake_Plan_H2_12",
    "Stocktake_Check_H2_13",
    "Stocktake_Summary_H2_14",
    "Impairment_H2_15",
    "Recoverable_H2_16",
    "Related_Party_H2_17",
]


def extract_sheet_structure(ws, sheet_name: str) -> dict:
    """提取单个sheet的完整结构信息"""
    result = {
        "sheet_name": sheet_name,
        "max_row": ws.max_row,
        "max_column": ws.max_column,
        "column_count": ws.max_column,
        "merged_ranges": [],
        "headers": [],
        "formula_cells": [],
        "data_types": {},
        "sample_data": [],
    }

    # 合并区域
    for merge in ws.merged_cells.ranges:
        result["merged_ranges"].append(str(merge))

    # 提取列头（前3行，通常为表头区域）
    header_rows = []
    for row_idx in range(1, min(4, ws.max_row + 1)):
        row_data = []
        for col_idx in range(1, ws.max_column + 1):
            cell = ws.cell(row=row_idx, column=col_idx)
            val = cell.value
            if val is not None:
                row_data.append({
                    "col": get_column_letter(col_idx),
                    "col_idx": col_idx,
                    "value": str(val)[:200],  # 截断避免太长
                })
        if row_data:
            header_rows.append({"row": row_idx, "cells": row_data})
    result["headers"] = header_rows

    # 扫描公式单元格和数据类型统计
    type_counts = {}
    formula_count = 0
    for row in ws.iter_rows(min_row=1, max_row=ws.max_row, max_col=ws.max_column):
        for cell in row:
            # 数据类型统计
            dt = cell.data_type
            type_counts[dt] = type_counts.get(dt, 0) + 1

            # 公式提取
            if cell.data_type == "f" or (
                isinstance(cell.value, str) and cell.value.startswith("=")
            ):
                formula_count += 1
                # 只记录前100个公式避免JSON过大
                if len(result["formula_cells"]) < 100:
                    result["formula_cells"].append({
                        "cell": f"{get_column_letter(cell.column)}{cell.row}",
                        "formula": str(cell.value)[:300],
                    })

    result["data_types"] = type_counts
    result["formula_count"] = formula_count

    # 样本数据（取前5行非空数据行，用于了解数据模式）
    data_start_row = 4  # 通常数据从第4行开始
    sample_rows_collected = 0
    for row_idx in range(data_start_row, min(data_start_row + 10, ws.max_row + 1)):
        row_data = []
        has_data = False
        for col_idx in range(1, min(ws.max_column + 1, 20)):  # 限制列数
            cell = ws.cell(row=row_idx, column=col_idx)
            val = cell.value
            if val is not None:
                has_data = True
                row_data.append({
                    "col": get_column_letter(col_idx),
                    "value": str(val)[:100],
                    "type": cell.data_type,
                })
        if has_data and sample_rows_collected < 5:
            result["sample_data"].append({"row": row_idx, "cells": row_data})
            sample_rows_collected += 1

    return result


def main():
    print(f"Reading: {XLSX_PATH}")
    if not XLSX_PATH.exists():
        print(f"ERROR: File not found: {XLSX_PATH}")
        return

    # data_only=False 保留公式；read_only=False 读取合并区域
    wb = openpyxl.load_workbook(str(XLSX_PATH), data_only=False, read_only=False)

    sheet_names = wb.sheetnames
    print(f"Total sheets found: {len(sheet_names)}")
    print(f"Sheet names: {sheet_names}")

    summary = {
        "source_file": str(XLSX_PATH),
        "total_sheets": len(sheet_names),
        "sheet_names": sheet_names,
        "expected_sheet_count": 21,
        "sheets_match_expected": len(sheet_names) == 21,
        "sheets": [],
    }

    for idx, name in enumerate(sheet_names):
        print(f"  [{idx+1}/{len(sheet_names)}] Processing: {name}")
        ws = wb[name]
        sheet_info = extract_sheet_structure(ws, name)
        summary["sheets"].append(sheet_info)

    # 生成映射表（spec glossary name → actual sheet name）
    print("\n--- Sheet Name Mapping (spec → actual) ---")
    summary["sheet_name_mapping"] = {}
    for i, actual_name in enumerate(sheet_names):
        expected_name = EXPECTED_SHEETS[i] if i < len(EXPECTED_SHEETS) else "UNEXPECTED"
        summary["sheet_name_mapping"][expected_name] = actual_name
        print(f"  {expected_name:40s} → {actual_name}")

    # 验证摘要
    summary["validation"] = {
        "sheet_count_ok": len(sheet_names) == 21,
        "all_sheets_have_data": all(
            s["max_row"] > 0 for s in summary["sheets"]
        ),
        "total_formula_cells": sum(
            s["formula_count"] for s in summary["sheets"]
        ),
        "total_merged_ranges": sum(
            len(s["merged_ranges"]) for s in summary["sheets"]
        ),
    }

    # 写出JSON
    OUTPUT_PATH.parent.mkdir(parents=True, exist_ok=True)
    with open(OUTPUT_PATH, "w", encoding="utf-8") as f:
        json.dump(summary, f, ensure_ascii=False, indent=2)

    print(f"\nOutput written to: {OUTPUT_PATH}")
    print(f"Validation: sheet_count={'OK' if summary['validation']['sheet_count_ok'] else 'MISMATCH'}")
    print(f"  Total formulas: {summary['validation']['total_formula_cells']}")
    print(f"  Total merged ranges: {summary['validation']['total_merged_ranges']}")

    # 逐sheet摘要
    print("\n--- Per-Sheet Summary ---")
    print(f"{'Sheet':<40s} {'Rows':>5s} {'Cols':>5s} {'Formulas':>8s} {'Merges':>7s}")
    print("-" * 70)
    for s in summary["sheets"]:
        print(
            f"{s['sheet_name']:<40s} {s['max_row']:>5d} {s['max_column']:>5d} "
            f"{s['formula_count']:>8d} {len(s['merged_ranges']):>7d}"
        )

    wb.close()


if __name__ == "__main__":
    main()

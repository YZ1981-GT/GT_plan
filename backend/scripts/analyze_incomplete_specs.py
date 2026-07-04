#!/usr/bin/env python3
"""Cross-reference incomplete active specs with code + xlsx templates."""
from __future__ import annotations

import json
import re
from pathlib import Path

ROOT = Path(__file__).resolve().parents[2]
SPECS = ROOT / ".kiro" / "specs"
WP_TEMPLATES = ROOT / "backend" / "wp_templates"
FRONTEND_WP = ROOT / "audit-platform" / "frontend" / "src" / "components" / "workpaper"
BACKEND = ROOT / "backend" / "app"

# Incomplete specs from prior scan (pending required tasks > 0 OR 0% done)
INCOMPLETE_SPECS = [
    "cutoff-test-auto-sampling", "d2-accounts-receivable-refactor", "d3-prepaid-accounts",
    "d7-contract-liabilities", "f0-confirmation", "display-format-single-source",
    "voucher-sampling-engine",
    "f1-prepayment", "f2-inventory-main", "f2-inventory-special", "f2-inventory-valuation-impairment",
    "f3-notes-payable", "f4-accounts-payable", "f5-cost-of-sales",
    "g0-confirmation", "g1-trading-financial-assets", "g2-interest-receivable",
    "g3-dividend-receivable", "g4-bond-investment-main", "g4-bond-investment-sppi",
    "g4-bond-investment-ecl", "g5-long-term-receivable", "g6-other-bond-investment-main",
    "g6-other-bond-investment-sppi", "g6-other-bond-investment-ecl",
    "g7-long-term-equity-main", "g7-long-term-equity-method", "g7-long-term-equity-subsidiary",
    "g8-other-equity-instruments", "g9-other-noncurrent-financial",
    "g10-trading-financial-liabilities", "g11-investment-income", "g12-net-hedge-gains",
    "g13-fair-value-changes", "g14-credit-impairment-loss",
]

SPEC_TO_COMPONENT = {
    "d2-accounts-receivable-refactor": "d2-accounts-receivable",
    "d3-prepaid-accounts": "d3-prepaid-accounts",
    "d7-contract-liabilities": "d7-contract-liabilities",
    "e1-monetary-fund-refactor": "e1-monetary-fund",
    "f1-prepayment": "f1-prepayment",
    "f2-inventory-main": "f2-inventory-main",
    "f3-notes-payable": "f3-notes-payable",
    "f4-accounts-payable": "f4-accounts-payable",
    "f5-cost-of-sales": "f5-cost-of-sales",
    "g1-trading-financial-assets": "g1-trading-financial-assets",
}


def parse_tasks(spec_dir: Path) -> dict:
    tasks_file = spec_dir / "tasks.md"
    if not tasks_file.exists():
        return {"done": 0, "pending": 0, "opt_pending": 0}
    text = tasks_file.read_text(encoding="utf-8")
    return {
        "done": len(re.findall(r"- \[x\]", text)),
        "pending": len(re.findall(r"- \[ \] ", text)),
        "opt_pending": len(re.findall(r"- \[ \]\*", text)),
    }


def grep_component(component: str) -> dict:
    hits = {"registry": False, "vue_main": [], "composables": [], "backend": []}
    reg = FRONTEND_WP / "htmlRendererRegistry.ts"
    if reg.exists() and component in reg.read_text(encoding="utf-8"):
        hits["registry"] = True
    for p in FRONTEND_WP.rglob("*"):
        if not p.is_file():
            continue
        name = p.name.lower()
        comp_short = component.replace("-", "")
        if component in p.read_text(encoding="utf-8", errors="ignore") or comp_short in name:
            rel = str(p.relative_to(ROOT))
            if p.suffix == ".vue" and p.name.startswith("Gt"):
                hits["vue_main"].append(rel)
            elif "composable" in str(p) or p.name.startswith("use"):
                hits["composables"].append(rel)
    for p in BACKEND.rglob("*.py"):
        if component in p.read_text(encoding="utf-8", errors="ignore"):
            hits["backend"].append(str(p.relative_to(ROOT)))
    return hits


def read_xlsx_summary(xlsx_path: Path) -> dict | None:
    try:
        import openpyxl
    except ImportError:
        return {"error": "openpyxl not installed"}
    if not xlsx_path.exists():
        return None
    wb = openpyxl.load_workbook(xlsx_path, read_only=True, data_only=True)
    sheets = []
    for ws in wb.worksheets:
        max_row = ws.max_row or 0
        max_col = ws.max_column or 0
        # sample header row (first non-empty row in first 5 rows)
        header = []
        for r in range(1, min(6, max_row + 1)):
            row_vals = [str(c.value or "").strip() for c in ws[r][:min(8, max_col)]]
            if any(row_vals):
                header = [v for v in row_vals if v][:6]
                break
        sheets.append({
            "name": ws.title,
            "rows": max_row,
            "cols": max_col,
            "header_sample": header,
        })
    wb.close()
    return {"file": xlsx_path.name, "sheet_count": len(sheets), "sheets": sheets}


def find_xlsx_for_spec(spec_name: str) -> list[Path]:
    """Map spec to wp_templates xlsx files."""
    mapping = {
        "f1-prepayment": ["F/F1 预付账款.xlsx"],
        "f2-inventory-main": ["F/F2-1至F2-14 存货及跌价准备-审定明细表类（Leap-常规程序）.xlsx"],
        "f3-notes-payable": ["F/F3 应付票据.xlsx"],
        "f4-accounts-payable": ["F/F4 应付账款.xlsx"],
        "f5-cost-of-sales": ["F/F5 营业成本.xlsx"],
        "g0-confirmation": ["G/G0 函证.xlsx"],
        "g1-trading-financial-assets": ["G/G1 交易性金融资产.xlsx"],
        "g2-interest-receivable": ["G/G2 应收利息.xlsx"],
        "g4-bond-investment-main": ["G/G4 债权投资.xlsx"],
        "g7-long-term-equity-main": ["G/G7 长期股权投资.xlsx"],
        "d2-accounts-receivable-refactor": ["D/D2-1至D2-4  应收账款- 审定表明细表（Leap-常规程序）.xlsx"],
        "d3-prepaid-accounts": ["D/D3 预收账款.xlsx"],
        "d7-contract-liabilities": ["D/D7 合同负债.xlsx"],
        "f0-confirmation": ["F/F0 函证.xlsx"],
    }
    paths = []
    for rel in mapping.get(spec_name, []):
        p = WP_TEMPLATES / rel
        if p.exists():
            paths.append(p)
    # fallback: glob by prefix
    if not paths:
        prefix = spec_name.split("-")[0].upper()
        if len(prefix) <= 2:
            cycle = prefix[0].upper()
            for p in sorted((WP_TEMPLATES / cycle).glob("*.xlsx"))[:1]:
                paths.append(p)
    return paths


def main():
    registry = json.loads((ROOT / "backend" / "data" / "account_package_registry.json").read_text(encoding="utf-8"))
    pkg_by_spec = {
        "d2-accounts-receivable-refactor": "D2_accounts_receivable",
        "d3-prepaid-accounts": "D3_prepaid_accounts",
        "d7-contract-liabilities": "D7_contract_liabilities",
        "f1-prepayment": "F1_prepayments",
        "f2-inventory-main": "F2_inventory",
        "f3-notes-payable": "F3_notes_payable",
        "f4-accounts-payable": "F4_accounts_payable",
        "f5-cost-of-sales": "F5_operating_cost",
        "g1-trading-financial-assets": "G1_trading_financial_assets",
    }
    packages = {p["account_package_id"]: p for p in registry["packages"]}

    results = []
    for spec in INCOMPLETE_SPECS:
        spec_dir = SPECS / spec
        if not spec_dir.exists():
            continue
        tasks = parse_tasks(spec_dir)
        component = SPEC_TO_COMPONENT.get(spec)
        code = grep_component(component) if component else {}
        pkg_id = pkg_by_spec.get(spec)
        pkg = packages.get(pkg_id) if pkg_id else None
        xlsx_paths = find_xlsx_for_spec(spec)
        xlsx_info = [read_xlsx_summary(p) for p in xlsx_paths[:1]]
        results.append({
            "spec": spec,
            "tasks": tasks,
            "componentType": component,
            "registered": code.get("registry", False),
            "vue_main_count": len(code.get("vue_main", [])),
            "vue_main": code.get("vue_main", [])[:3],
            "composable_count": len(code.get("composables", [])),
            "backend_hits": len(code.get("backend", [])),
            "registry_sheets": len(pkg["sheets"]) if pkg else None,
            "xlsx": xlsx_info[0] if xlsx_info else None,
        })

    print(json.dumps(results, ensure_ascii=False, indent=2))


if __name__ == "__main__":
    main()

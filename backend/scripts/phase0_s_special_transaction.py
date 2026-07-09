"""Phase0: S-special-transaction-workpapers 双源核对 + S17 格式转换.

Deliverables:
1. 确认 14 底稿 sheet 结构（审定表/内控调查表/判断子表/专家多分支）
2. S17 .xls → .xlsx 转换（xlrd 读取 + openpyxl 写入）
3. sheetName 分发表
4. 引擎输入/输出规格

Usage:
    python backend/scripts/phase0_s_special_transaction.py
"""
from __future__ import annotations

import json
import sys
from pathlib import Path
from typing import Any

import openpyxl
import xlrd

# ─── Paths ───────────────────────────────────────────────────────────────────
ROOT = Path(
    r"d:\GT_plan\数据\致同通用审计程序及底稿模板（2025年修订）"
    r"\1.致同审计程序及底稿模板（2025年）\6.特定项目程序（S）"
)
OUTPUT_DIR = Path(r"d:\GT_plan\backend\scripts\output")
OUTPUT_DIR.mkdir(parents=True, exist_ok=True)

# ─── Target files (14 workpapers for this spec) ─────────────────────────────
TARGETS: dict[str, str] = {
    "S1": "S1 对被审计单位违反法规行为的考虑2019.xlsx",
    "S2": "S2 首次接受委托时对期初余额的审计2019.xlsx",
    "S4": "S4 非货币性资产交换202401.xlsx",
    "S5": "S5 债务重组202401.xlsx",
    "S6": "S6 对大股东及关联方资金占用和违规担保情况审核程序.xlsx",
    "S8": "S8 租赁.xlsx",
    "S9": "S9 对电子商务的考虑.xlsx",
    "S10": "S10 对环境事项的考虑.xlsx",
    "S11": "S11 利用服务机构活动.xlsx",
    "S12": "S12 利用专家（注册会计师的专家）的工作.xlsx",
    "S13": "S13 利用管理层的专家编制的信息形成审计证据.xlsx",
    "S14": "S14 会计估计和相关披露.xlsx",
    "S16": "S16 套期活动.xlsx",
    "S17": "S17 非经常性损益.xls",
}


def analyze_xlsx(wp_code: str, filepath: Path) -> dict[str, Any]:
    """Analyze an xlsx file and return sheet structure."""
    wb = openpyxl.load_workbook(filepath, read_only=False, data_only=False)
    sheets = []
    for name in wb.sheetnames:
        ws = wb[name]
        rows = ws.max_row or 0
        cols = ws.max_column or 0
        merged = len(ws.merged_cells.ranges)
        formulas = 0
        for row in ws.iter_rows():
            for cell in row:
                if isinstance(cell.value, str) and cell.value.startswith("="):
                    formulas += 1
        sheets.append({
            "name": name,
            "rows": rows,
            "cols": cols,
            "merged": merged,
            "formulas": formulas,
        })
    wb.close()
    return {"wp_code": wp_code, "file": filepath.name, "sheets": sheets}


def analyze_xls_with_xlrd(wp_code: str, filepath: Path) -> dict[str, Any]:
    """Analyze an .xls file using xlrd and return sheet structure."""
    wb = xlrd.open_workbook(str(filepath), formatting_info=False)
    sheets = []
    for name in wb.sheet_names():
        ws = wb.sheet_by_name(name)
        sheets.append({
            "name": name,
            "rows": ws.nrows,
            "cols": ws.ncols,
            "merged": len(ws.merged_cells),
            "formulas": 0,  # xlrd 2.x doesn't expose formulas easily
        })
    return {"wp_code": wp_code, "file": filepath.name, "sheets": sheets}


def convert_s17_xls_to_xlsx(src: Path, dst: Path) -> None:
    """Convert S17 .xls to .xlsx using xlrd + openpyxl."""
    xls_wb = xlrd.open_workbook(str(src), formatting_info=False)
    xlsx_wb = openpyxl.Workbook()
    # Remove default sheet
    xlsx_wb.remove(xlsx_wb.active)

    for sheet_name in xls_wb.sheet_names():
        xls_ws = xls_wb.sheet_by_name(sheet_name)
        xlsx_ws = xlsx_wb.create_sheet(title=sheet_name)
        for row_idx in range(xls_ws.nrows):
            for col_idx in range(xls_ws.ncols):
                cell = xls_ws.cell(row_idx, col_idx)
                value = cell.value
                # Handle xlrd date cells
                if cell.ctype == xlrd.XL_CELL_DATE:
                    try:
                        value = xlrd.xldate_as_datetime(value, xls_wb.datemode)
                    except Exception:
                        pass
                xlsx_ws.cell(row=row_idx + 1, column=col_idx + 1, value=value)

        # Copy merged cells
        for crange in xls_ws.merged_cells:
            rlo, rhi, clo, chi = crange
            xlsx_ws.merge_cells(
                start_row=rlo + 1, start_column=clo + 1,
                end_row=rhi, end_column=chi
            )

    xlsx_wb.save(str(dst))
    xlsx_wb.close()
    print(f"  ✅ S17 converted: {src.name} → {dst.name} ({dst.stat().st_size:,} bytes)")


def classify_sheet(name: str) -> str:
    """Classify a sheet by its name pattern."""
    lower = name.lower()
    if "表头" in name:
        return "header"
    if "gt_custom" in lower:
        return "gt_custom"
    if "审定表" in name:
        return "adjudication"
    if "内控" in name or "内部控制" in name:
        return "internal_control"
    if "商业实质" in name:
        return "judgment_commercial_substance"
    if "损益确认时点" in name:
        return "judgment_timing"
    if "监管" in name or "风险提示" in name:
        return "regulatory_reference"
    if "程序表" in name or ("程序" in name and ("S" in name or "s" in name)):
        return "program"
    if "法律法规" in name or "环境法规" in name:
        return "regulation"
    if "说明" in name:
        return "description"
    if "沟通" in name:
        return "communication"
    if "记录" in name:
        return "record"
    if "评价" in name or "评估" in name:
        return "evaluation"
    if "了解" in name:
        return "understanding"
    if "应对" in name:
        return "response"
    if "偏向" in name:
        return "bias_indication"
    if "报告" in name:
        return "report"
    if "非经常性" in name:
        return "non_recurring"
    if "审核表" in name:
        return "review"
    if "核对表" in name:
        return "reconciliation"
    if "所得税影响" in name:
        return "tax_impact"
    return "other"


def build_dispatch_table(analyses: list[dict]) -> dict[str, Any]:
    """Build the sheetName dispatch table for Vue components."""
    dispatch = {}
    for item in analyses:
        wp_code = item["wp_code"]
        sheets_info = []
        for s in item["sheets"]:
            name = s["name"]
            classification = classify_sheet(name)
            if classification in ("header", "gt_custom"):
                continue  # Skip header/internal sheets
            sheets_info.append({
                "sheetName": name,
                "type": classification,
                "rows": s["rows"],
                "cols": s["cols"],
                "formulas": s["formulas"],
            })
        dispatch[wp_code] = {
            "file": item["file"],
            "componentType": _get_component_type(wp_code),
            "sheets": sheets_info,
        }
    return dispatch


def _get_component_type(wp_code: str) -> str:
    """Determine componentType per design doc."""
    mapping = {
        "S4": "s4-nonmonetary-exchange",
        "S5": "s5-debt-restructuring",
        "S6": "s6-fund-occupation",
        "S12": "s12-cpa-expert",
        "S13": "s13-mgmt-expert",
        "S14": "s14-accounting-estimate",
    }
    return mapping.get(wp_code, "a-program-console")


def build_engine_spec() -> dict[str, Any]:
    """Build engine input/output specifications per design doc."""
    return {
        "useS4FormulaEngine": {
            "functions": {
                "calcExchangeGainLoss": {
                    "input": {
                        "inFairValue": "number (换入资产公允价值)",
                        "outFairValue": "number (换出资产公允价值)",
                        "outBookValue": "number (换出资产账面价值)",
                        "taxes": "number (相关税费)",
                    },
                    "output": {
                        "gainLoss": "number (换出资产损益 = outFairValue - outBookValue)",
                        "inCost": "number (换入资产成本 = outFairValue + taxes)",
                    },
                },
                "judgeApplicable": {
                    "input": {
                        "exclusions": "boolean[6] (6项排除情形, true=属于该情形)",
                        "cashflowDifferent": "boolean (现金流量风险/时间/金额显著不同)",
                    },
                    "output": "boolean (是否适用非货币性资产交换准则: 所有exclusions为false→适用)",
                },
                "judgeCommercialSubstance": {
                    "input": {
                        "exclusions": "boolean[6]",
                        "cashflowDifferent": "boolean",
                    },
                    "output": "boolean (是否具有商业实质 = cashflowDifferent)",
                },
            },
        },
        "useS5FormulaEngine": {
            "functions": {
                "calcCreditorGainLoss": {
                    "input": {
                        "origBook": "number (债权原账面价值)",
                        "origFair": "number (债权原公允价值)",
                        "recvFair": "number (收到资产公允价值)",
                        "otherCost": "number (其他相关费用)",
                    },
                    "output": "number (债权人重组损益)",
                },
                "calcDebtorGainLoss": {
                    "input": {
                        "debtBook": "number (所清偿债务账面价值)",
                        "assetBook": "number (转让资产账面价值)",
                        "equityFair": "number (权益工具公允价值)",
                    },
                    "output": "number (债务人重组损益 = debtBook - assetBook - equityFair)",
                },
            },
        },
        "resolveExpertSubSheet": {
            "input": "ExpertDomain: 'general' | 'share-based-payment' | 'financial-instrument-fair-value'",
            "output_mapping": {
                "S12": {
                    "general": "S12-3-2 利用专家评价管理层的工作的适当性",
                    "share-based-payment": "S12-3-3 利用专家评价管理层的工作(股份支付)",
                    "financial-instrument-fair-value": "S12-3-4 利用专家评价管理层的工作(金融工具公允价值)",
                },
                "S13": {
                    "general": "S13-3-2 评价管理层专家工作的适当性",
                    "share-based-payment": "S13-3-3 评价管理层专家工作的适当性(股份支付)",
                    "financial-instrument-fair-value": "S13-3-4 评价管理层专家工作的适当性(金融工具公允价值)",
                },
            },
        },
    }


def main() -> None:
    sys.stdout.reconfigure(encoding="utf-8")
    print("=" * 72)
    print("Phase0: S-special-transaction-workpapers 双源核对 + S17 格式转换")
    print("=" * 72)

    # ─── Step 1: Analyze all 14 workpapers ───────────────────────────────
    print("\n## Step 1: 分析 14 底稿 sheet 结构\n")
    analyses: list[dict] = []
    errors: list[str] = []

    for wp_code, filename in sorted(TARGETS.items(), key=lambda x: x[0]):
        filepath = ROOT / filename
        if not filepath.exists():
            msg = f"  ❌ {wp_code}: 文件不存在 → {filepath}"
            print(msg)
            errors.append(msg)
            continue

        if filepath.suffix.lower() == ".xls":
            # S17: use xlrd
            print(f"  📄 {wp_code}: {filename} (.xls → xlrd)")
            try:
                result = analyze_xls_with_xlrd(wp_code, filepath)
                analyses.append(result)
                for s in result["sheets"]:
                    print(f"      [{s['name']}] {s['rows']}x{s['cols']} merged={s['merged']}")
            except Exception as e:
                msg = f"  ❌ {wp_code}: xlrd 读取失败 → {e}"
                print(msg)
                errors.append(msg)
        else:
            # xlsx: use openpyxl
            print(f"  📄 {wp_code}: {filename}")
            try:
                result = analyze_xlsx(wp_code, filepath)
                analyses.append(result)
                for s in result["sheets"]:
                    print(f"      [{s['name']}] {s['rows']}x{s['cols']} merged={s['merged']} formula={s['formulas']}")
            except Exception as e:
                msg = f"  ❌ {wp_code}: openpyxl 读取失败 → {e}"
                print(msg)
                errors.append(msg)

    # ─── Step 2: Convert S17 .xls → .xlsx ────────────────────────────────
    print("\n## Step 2: S17 .xls → .xlsx 转换\n")
    s17_src = ROOT / TARGETS["S17"]
    s17_dst = OUTPUT_DIR / "S17 非经常性损益.xlsx"

    if not s17_src.exists():
        msg = f"  ❌ S17 源文件不存在: {s17_src}"
        print(msg)
        errors.append(msg)
    else:
        try:
            convert_s17_xls_to_xlsx(s17_src, s17_dst)
        except Exception as e:
            msg = f"  ❌ S17 转换失败: {e}"
            print(msg)
            errors.append(msg)
            # Per Req 8.2: 不静默兜底空数据
            raise RuntimeError(f"S17 .xls conversion failed: {e}") from e

    # ─── Step 3: Build sheetName dispatch table ──────────────────────────
    print("\n## Step 3: sheetName 分发表\n")
    dispatch = build_dispatch_table(analyses)
    for wp_code, info in sorted(dispatch.items()):
        ct = info["componentType"]
        print(f"  {wp_code} → {ct}")
        for s in info["sheets"]:
            print(f"      [{s['sheetName']}] type={s['type']} {s['rows']}x{s['cols']}")

    # ─── Step 4: Engine input/output spec ────────────────────────────────
    print("\n## Step 4: 引擎输入/输出规格\n")
    engine_spec = build_engine_spec()
    for engine_name, spec in engine_spec.items():
        print(f"  {engine_name}:")
        if "functions" in spec:
            for fn_name, fn_spec in spec["functions"].items():
                print(f"    {fn_name}:")
                if isinstance(fn_spec.get("input"), dict):
                    for k, v in fn_spec["input"].items():
                        print(f"      in: {k} → {v}")
                else:
                    print(f"      in: {fn_spec.get('input')}")
                out = fn_spec.get("output")
                if isinstance(out, dict):
                    for k, v in out.items():
                        print(f"      out: {k} → {v}")
                else:
                    print(f"      out: {out}")
        elif "output_mapping" in spec:
            print(f"    input: {spec['input']}")
            for prefix, mapping in spec["output_mapping"].items():
                for domain, sheet_name in mapping.items():
                    print(f"    {prefix}.{domain} → {sheet_name}")

    # ─── Step 5: Save JSON outputs ───────────────────────────────────────
    output_data = {
        "workpapers_analyzed": len(analyses),
        "errors": errors,
        "dispatch_table": dispatch,
        "engine_spec": engine_spec,
    }
    output_file = OUTPUT_DIR / "s_special_transaction_phase0.json"
    with open(output_file, "w", encoding="utf-8") as f:
        json.dump(output_data, f, ensure_ascii=False, indent=2)
    print(f"\n## Output saved: {output_file}")

    # ─── Summary ─────────────────────────────────────────────────────────
    print("\n" + "=" * 72)
    print("Summary:")
    print(f"  ✅ 分析完成: {len(analyses)}/14 底稿")
    if errors:
        print(f"  ❌ 错误: {len(errors)} 条")
        for e in errors:
            print(f"     {e}")
    else:
        print("  ✅ 无错误")

    # Verify S17 conversion
    if s17_dst.exists():
        print(f"  ✅ S17 已转换: {s17_dst} ({s17_dst.stat().st_size:,} bytes)")
    print("=" * 72)


if __name__ == "__main__":
    main()

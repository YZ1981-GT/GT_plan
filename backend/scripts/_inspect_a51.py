"""Inspect A5-1 template structure."""
import openpyxl

wb = openpyxl.load_workbook(r"wp_templates/A/A5-1 现金流量表审计.xlsx", read_only=True, data_only=True)
print("Sheets:", wb.sheetnames)
for sn in wb.sheetnames:
    ws = wb[sn]
    print(f"\n{'='*60}")
    print(f"Sheet: {sn} (rows={ws.max_row}, cols={ws.max_column})")
    print("="*60)
    for i, row in enumerate(ws.iter_rows(min_row=1, max_row=min(50, ws.max_row or 1), values_only=True), 1):
        cells = [str(c)[:60] if c else "" for c in row[:12]]
        line = " | ".join(cells)
        if line.strip(" |"):
            print(f"  R{i:02d}: {line}")
wb.close()

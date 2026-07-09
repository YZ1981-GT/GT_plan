"""
analyze_h3_investment_property.py
=================================
读取 H3 投资性房地产.xlsx 全部 sheet，提取结构信息并输出 JSON。

Phase 0 双源输入验证 - 任务 0.1

提取：sheet名 / 列头 / 行数 / 公式单元格 / 合并区域 / 数据类型
产出：.kiro/specs/h3-investment-property/h3_structure_summary.json
验证：22 sheet 结构与 spec design.md 描述一致

用法：python backend/scripts/analyze_h3_investment_property.py
"""

import json
import os
import sys
from pathlib import Path
from collections import Counter

import openpyxl
from openpyxl.utils import get_column_letter


# ---------- 配置 ----------

_CANDIDATES = [
    Path(
        r"基础数据/致同通用审计程序及底稿模板（2025年修订）"
        r"/1.致同审计程序及底稿模板（2025年）"
        r"/4.风险应对-实质性程序（D-N）/H 固定资产循环/H3 投资性房地产.xlsx"
    ),
]

OUTPUT_FILE = Path(".kiro/specs/h3-investment-property/h3_structure_summary.json")

# 根据 design.md 描述的 22 个 sheet 预期结构
# H3-1(成本/公允双版本) H3-2(成本/公允双版本) H3-3 H3-4 H3-5(成本/公允双版本)
# H3-6 H3-7(不含减值/含减值双分支) H3-8 H3-9 H3-10 H3-11 H3-12 H3-13 H3-14
# H3A(程序表) 底稿目录 附注上市 附注国企
EXPECTED_SHEETS = {
    "Tab_Index": {"keyword": "目录", "expected_cols_approx": 8},
    "Procedure_Table_H3A": {"keyword": "H3A", "expected_cols_approx": 13},
    "Adjudication_H3_1_Cost": {"keyword": "审定", "expected_cols_approx": 10},
    "Adjudication_H3_1_Fair": {"keyword": "公允", "note": "H3-1公允价值版本"},
    "Detail_H3_2_Cost": {"keyword": "明细", "expected_cols_approx": 49},
    "Detail_H3_2_Fair": {"keyword": "公允", "note": "H3-2公允价值版本"},
    "Adjustment_H3_3": {"keyword": "H3-3", "expected_cols_approx": 10},
    "Policy_Check_H3_4": {"keyword": "H3-4", "expected_cols_approx": 16},
    "Addition_Check_H3_5_Cost": {"keyword": "H3-5", "note": "成本模式增减检查"},
    "Addition_Check_H3_5_Fair": {"keyword": "H3-5", "note": "公允模式增减检查"},
    "Transfer_Review_H3_6": {"keyword": "H3-6", "expected_cols_approx": 36},
    "Depreciation_NoImpair_H3_7": {"keyword": "不含减值", "expected_cols_approx": 28},
    "Depreciation_WithImpair_H3_7": {"keyword": "含减值", "expected_cols_approx": 28},
    "Fair_Value_Review_H3_8": {"keyword": "H3-8", "expected_cols_approx": 15},
    "Stocktake_H3_9": {"keyword": "H3-9", "expected_cols_approx": 13},
    "Impairment_H3_10": {"keyword": "H3-10", "expected_cols_approx": 20},
    "Recoverable_H3_11": {"keyword": "H3-11", "expected_cols_approx": 28},
    "Title_Check_H3_12": {"keyword": "H3-12", "expected_cols_approx": 16},
    "Related_Party_H3_13": {"keyword": "H3-13", "expected_cols_approx": 11},
    "Rental_Income_H3_14": {"keyword": "H3-14", "expected_cols_approx": 27},
    "Disclosure_Listed": {"keyword": "上市", "expected_cols_approx": 100},
    "Disclosure_SOE": {"keyword": "国有", "expected_cols_approx": 100},
}


def find_source_file() -> Path:
    """按候选列表查找第一个存在的 xlsx 文件。"""
    for candidate in _CANDIDATES:
        if candidate.exists():
            return candidate
    # 尝试绝对路径（从项目根）
    project_root = Path(os.getcwd())
    for candidate in _CANDIDATES:
        abs_path = project_root / candidate
        if abs_path.exists():
            return abs_path
    print("ERROR: 找不到 H3 投资性房地产.xlsx，尝试过：")
    for c in _CANDIDATES:
        print(f"  - {c}")
    sys.exit(1)


def extract_headers(ws, max_scan_rows: int = 5) -> dict:
    """从前几行提取列头信息，返回所有可能的列头行。"""
    header_rows = {}
    for row_idx in range(1, min(max_scan_rows + 1, ws.max_row + 1)):
        headers = []
        for col_idx in range(1, ws.max_column + 1):
            cell = ws.cell(row=row_idx, column=col_idx)
            val = cell.value
            if val is not None:
                headers.append(str(val).strip())
            else:
                headers.append("")
        non_empty = [h for h in headers if h]
        if len(non_empty) >= 2:
            header_rows[f"row_{row_idx}"] = non_empty
    return header_rows


def extract_primary_headers(ws, max_scan_rows: int = 5) -> list[str]:
    """提取最可能是主列头的一行（非空数最多的行）。"""
    best_row = []
    best_count = 0
    for row_idx in range(1, min(max_scan_rows + 1, ws.max_row + 1)):
        headers = []
        for col_idx in range(1, ws.max_column + 1):
            cell = ws.cell(row=row_idx, column=col_idx)
            val = cell.value
            if val is not None:
                headers.append(str(val).strip())
            else:
                headers.append("")
        non_empty = sum(1 for h in headers if h)
        if non_empty > best_count:
            best_count = non_empty
            best_row = [h for h in headers if h]
    return best_row


def extract_formulas(ws) -> dict:
    """提取公式单元格信息。"""
    formulas = []
    formula_count = 0
    formula_by_col = Counter()
    for row in ws.iter_rows(min_row=1, max_row=ws.max_row, max_col=ws.max_column):
        for cell in row:
            if cell.value and isinstance(cell.value, str) and cell.value.startswith("="):
                formula_count += 1
                col_letter = get_column_letter(cell.column)
                formula_by_col[col_letter] += 1
                if len(formulas) < 20:  # 保留前 20 个样例
                    formulas.append({
                        "cell": cell.coordinate,
                        "formula": cell.value,
                    })
    return {
        "count": formula_count,
        "samples": formulas,
        "by_column": dict(formula_by_col.most_common(10)),
    }


def extract_merged_areas(ws) -> list[str]:
    """提取合并单元格区域。"""
    return [str(r) for r in ws.merged_cells.ranges]


def extract_data_types(ws) -> dict:
    """统计各列数据类型分布（采样前 30 行数据区）。"""
    type_summary = {}
    data_start_row = min(4, ws.max_row)  # 跳过列头
    sample_end_row = min(data_start_row + 30, ws.max_row)

    for col_idx in range(1, ws.max_column + 1):
        col_letter = get_column_letter(col_idx)
        types = Counter()
        for row_idx in range(data_start_row, sample_end_row + 1):
            cell = ws.cell(row=row_idx, column=col_idx)
            val = cell.value
            if val is None:
                types["null"] += 1
            elif isinstance(val, str):
                if val.startswith("="):
                    types["formula"] += 1
                else:
                    types["string"] += 1
            elif isinstance(val, (int, float)):
                types["number"] += 1
            elif hasattr(val, "strftime"):
                types["date"] += 1
            else:
                types["other"] += 1

        # 只保留非空类型
        non_null_types = {k: v for k, v in types.items() if k != "null" and v > 0}
        if non_null_types:
            dominant_type = max(non_null_types, key=non_null_types.get)
            type_summary[col_letter] = dominant_type

    return type_summary


def validate_against_spec(sheets: list[dict]) -> dict:
    """将实际 sheet 结构与 spec design.md 中 22 sheet 预期进行对照验证。"""
    validation = {
        "total_sheets": len(sheets),
        "expected_sheets": 22,
        "match": len(sheets) == 22,
        "matched_glossary_entries": [],
        "unmatched_glossary_entries": [],
        "sheet_mapping": {},
    }

    sheet_names = [s["name"] for s in sheets]

    for glossary_key, expected in EXPECTED_SHEETS.items():
        keyword = expected["keyword"]
        matched_sheet = None
        for sname in sheet_names:
            if keyword in sname:
                # 避免重复匹配同一 sheet
                if sname not in validation["sheet_mapping"].values():
                    matched_sheet = sname
                    break
        # 如果第一次没匹配到，放宽搜索（允许已匹配的也再匹配）
        if not matched_sheet:
            for sname in sheet_names:
                if keyword in sname:
                    matched_sheet = sname
                    break

        if matched_sheet:
            validation["matched_glossary_entries"].append(glossary_key)
            validation["sheet_mapping"][glossary_key] = matched_sheet
        else:
            validation["unmatched_glossary_entries"].append(glossary_key)

    return validation


def classify_sheets_by_measurement_model(sheets: list[dict]) -> dict:
    """根据sheet名称分类：成本模式专属/公允模式专属/共用。"""
    classification = {
        "cost_only": [],
        "fair_only": [],
        "shared": [],
        "dual_version": [],
    }

    # 关键词判断
    cost_keywords = ["成本", "折旧", "减值", "不含减值", "含减值"]
    fair_keywords = ["公允价值"]

    for sheet in sheets:
        name = sheet["name"]
        if any(kw in name for kw in cost_keywords) and any(kw in name for kw in fair_keywords):
            classification["dual_version"].append(name)
        elif any(kw in name for kw in cost_keywords):
            classification["cost_only"].append(name)
        elif any(kw in name for kw in fair_keywords):
            classification["fair_only"].append(name)
        else:
            classification["shared"].append(name)

    return classification


def main():
    # 切换到项目根目录
    project_root = Path(__file__).resolve().parent.parent.parent
    os.chdir(project_root)

    source_file = find_source_file()
    print(f"📂 读取文件: {source_file}")
    print(f"   文件大小: {source_file.stat().st_size / 1024:.1f} KB")

    # data_only=False 以读取公式而非计算结果
    wb = openpyxl.load_workbook(str(source_file), data_only=False, read_only=False)

    result = {
        "file_path": str(source_file),
        "file_size_kb": round(source_file.stat().st_size / 1024, 1),
        "sheet_count": len(wb.sheetnames),
        "sheet_names": wb.sheetnames,
        "sheets": [],
        "measurement_model_classification": None,
        "validation": None,
        "summary": {},
    }

    print(f"\n共 {len(wb.sheetnames)} 个 sheet:")
    print("-" * 70)

    total_formulas = 0
    total_merged = 0

    for idx, name in enumerate(wb.sheetnames, 1):
        ws = wb[name]
        print(f"  [{idx:2d}] {name}")
        print(f"       尺寸: {ws.max_row}行 × {ws.max_column}列")

        # 提取各项结构信息
        all_headers = extract_headers(ws)
        primary_headers = extract_primary_headers(ws)
        formula_info = extract_formulas(ws)
        merged_areas = extract_merged_areas(ws)
        data_types = extract_data_types(ws)

        total_formulas += formula_info["count"]
        total_merged += len(merged_areas)

        print(f"       列头: {len(primary_headers)}列 | 公式: {formula_info['count']}个 | 合并: {len(merged_areas)}个")

        sheet_info = {
            "index": idx,
            "name": name,
            "rows": ws.max_row,
            "cols": ws.max_column,
            "dimensions": f"{ws.max_row}行 × {ws.max_column}列",
            "primary_headers": primary_headers,
            "all_header_rows": all_headers,
            "formula_count": formula_info["count"],
            "formula_samples": formula_info["samples"],
            "formula_by_column": formula_info["by_column"],
            "merged_areas": merged_areas,
            "merged_area_count": len(merged_areas),
            "data_types": data_types,
        }
        result["sheets"].append(sheet_info)

    # 计量模式分类
    print("\n" + "=" * 70)
    print("📊 计量模式分类:")
    classification = classify_sheets_by_measurement_model(result["sheets"])
    result["measurement_model_classification"] = classification
    print(f"  成本模式专属: {classification['cost_only']}")
    print(f"  公允模式专属: {classification['fair_only']}")
    print(f"  双版本: {classification['dual_version']}")
    print(f"  共用: {classification['shared']}")

    # 验证
    print("\n" + "=" * 70)
    print("📋 Spec design.md 对照验证:")
    validation = validate_against_spec(result["sheets"])
    result["validation"] = validation

    if validation["match"]:
        print(f"  ✅ sheet 数量匹配: {validation['total_sheets']}/{validation['expected_sheets']}")
    else:
        print(f"  ⚠️ sheet 数量不匹配: 实际 {validation['total_sheets']} vs 预期 {validation['expected_sheets']}")

    print(f"  已匹配预期条目: {len(validation['matched_glossary_entries'])}/{len(EXPECTED_SHEETS)}")
    if validation["unmatched_glossary_entries"]:
        print(f"  ❌ 未匹配条目: {validation['unmatched_glossary_entries']}")

    # 汇总
    result["summary"] = {
        "total_sheets": len(wb.sheetnames),
        "total_formulas": total_formulas,
        "total_merged_areas": total_merged,
        "total_rows": sum(s["rows"] for s in result["sheets"]),
        "total_cols_max": max(s["cols"] for s in result["sheets"]),
        "avg_rows_per_sheet": round(sum(s["rows"] for s in result["sheets"]) / len(result["sheets"]), 1),
        "formula_dense_sheets": sorted(
            [(s["name"], s["formula_count"]) for s in result["sheets"] if s["formula_count"] > 10],
            key=lambda x: x[1],
            reverse=True,
        ),
    }

    # 输出
    OUTPUT_FILE.parent.mkdir(parents=True, exist_ok=True)
    with open(OUTPUT_FILE, "w", encoding="utf-8") as f:
        json.dump(result, f, ensure_ascii=False, indent=2)

    print(f"\n✅ 产出文件: {OUTPUT_FILE}")
    print(f"   sheet 数量: {result['sheet_count']}")
    print(f"   公式总数: {total_formulas}")
    print(f"   合并区域总数: {total_merged}")
    print(f"   数据行总数: {result['summary']['total_rows']}")

    # 验证 22 sheet
    if result["sheet_count"] != 22:
        print(f"\n⚠️  注意：实际 sheet 数量为 {result['sheet_count']}，spec 描述为 22。")
        print("   可能存在隐藏 sheet 或 sheet 数量描述需要修正。")
    else:
        print(f"\n✅ sheet 数量验证通过: 22 sheet 与 spec 描述一致")

    wb.close()
    return result


if __name__ == "__main__":
    main()

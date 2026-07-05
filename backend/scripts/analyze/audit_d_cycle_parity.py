#!/usr/bin/env python3
"""Audit D1-D7 workpaper parity vs D4 benchmark + xlsx sheet coverage."""
from __future__ import annotations

import json
import re
import sys
from pathlib import Path
from collections import defaultdict

try:
    import openpyxl
except ImportError:
    print("ERROR: pip install openpyxl")
    sys.exit(1)

ROOT = Path(r"D:\GT_plan")
TEMPLATE_DIR = ROOT / "backend" / "wp_templates" / "D"
FRONTEND = ROOT / "audit-platform" / "frontend" / "src" / "components" / "workpaper"

# Map wp_code -> xlsx file patterns (exclude D0)
WP_FILES: dict[str, list[str]] = {
    "D1": ["D1 应收票据.xlsx"],
    "D2": [
        "D2-1至D2-4  应收账款- 审定表明细表（Leap-常规程序）.xlsx",
        "D2-5  应收账款 -分析程序（Leap应对措施-分析程序）.xlsx",
        "D2-6至D2-13  应收账款 -检查（Leap应对措施-检查）.xlsx",
    ],
    "D3": ["D3 预收账款.xlsx"],
    "D4": [
        "D4-1至D4-4 营业收入 - 审定表明细表（Leap-常规程序）.xlsx",
        "D4-5 营业收入-会计政策（Leap-常规程序）.xlsx",
        "D4-6至D4-11营业收入 - 分析程序（Leap应对措施-分析程序）.xlsx",
        "D4-12 营业收入-合同检查（Leap-常规程序）.xlsx",
        "D4-13至D4-20主营业务收入-检查（Leap应对措施-检查）.xlsx",
        "D4-21营业收入-关联方检查（Leap-常规程序）.xlsx",
        "D4-22至D4-32营业收入-IPO 上市 新三板 重组 舞弊应对.xlsx",
        "D4-33至D4-36 其他业务收入.xlsx",
    ],
    "D5": ["D5 应收款项融资.xlsx"],
    "D6": ["D6 合同资产.xlsx"],
    "D7": ["D7 合同负债.xlsx"],
}

SHEET_CODE_RE = re.compile(
    r"(D[1-7](?:-\d+[A-Z]?|A|A原)?|附注(?:上市|国企)|GT_Custom|目录|底稿目录|业务模式分析提示|访谈记录与核对示例|修订前程序表|D7A原|D8A原)",
    re.I,
)


def extract_code_from_sheet(name: str) -> str | None:
    """Best-effort sheet code from xlsx tab name."""
    n = name.strip()
    if n in ("目录", "底稿目录", "D4", "D1", "D2", "D3", "D5", "D6", "D7"):
        return "directory"
    if "附注" in n and "上市" in n:
        return "附注上市"
    if "附注" in n and ("国企" in n or "国有" in n):
        return "附注国企"
    if "访谈记录与核对" in n:
        return "D4-31T"
    if "业务模式分析提示" in n:
        return "业务模式提示"
    if "GT_Custom" in n or n == "GT_Custom":
        return "GT_Custom"
    if "修订前程序表" in n:
        return "修订前程序表"
    # D6/D7 legacy program sheets
    if re.search(r"D7A原", n):
        return "D7A原"
    if re.search(r"D8A原", n):
        return "D8A原"
    m = re.search(r"(D[1-7]A)\b", n)
    if m:
        return m.group(1)
    m = re.search(r"(D[1-7]-\d+[A-Z]?)\s*$", n)
    if m:
        return m.group(1)
    m = re.search(r"(D[1-7]-\d+[A-Z]?)", n)
    if m:
        return m.group(1)
    if "审定表" in n and re.search(r"D5", n):
        return "D5-1"  # special: 审定表D5
    return None


def read_xlsx_sheets(wp: str) -> list[dict]:
    rows = []
    for fname in WP_FILES[wp]:
        path = TEMPLATE_DIR / fname
        if not path.exists():
            # fuzzy match
            hits = list(TEMPLATE_DIR.glob(fname.split()[0] + "*"))
            if not hits:
                rows.append({"file": fname, "error": "NOT FOUND"})
                continue
            path = hits[0]
        wb = openpyxl.load_workbook(path, read_only=True, data_only=True)
        for sn in wb.sheetnames:
            ws = wb[sn]
            code = extract_code_from_sheet(sn)
            rows.append({
                "file": path.name,
                "sheet_name": sn,
                "code": code,
                "max_row": ws.max_row or 0,
                "max_col": ws.max_column or 0,
            })
        wb.close()
    return rows


def parse_index_rows_ts(path: Path) -> list[dict]:
    """Parse INDEX_ROWS from d*SheetLabels.ts or D2TabIndex.vue."""
    if not path.exists():
        return []
    text = path.read_text(encoding="utf-8")
    rows = []
    for m in re.finditer(
        r"\{\s*seq:\s*(\d+),\s*name:\s*'([^']+)',\s*code:\s*'([^']+)',\s*group:\s*'([^']+)',\s*tabKey:\s*'([^']+)',\s*applicable:\s*(true|false)",
        text,
    ):
        rows.append({
            "seq": int(m.group(1)),
            "name": m.group(2),
            "code": m.group(3),
            "group": m.group(4),
            "tabKey": m.group(5),
            "applicable": m.group(6) == "true",
        })
    return rows


def parse_gt_routes(path: Path) -> set[str]:
    if not path.exists():
        return set()
    text = path.read_text(encoding="utf-8")
    codes = set()
    for m in re.finditer(r"currentSheet === '([^']+)'", text):
        codes.add(m.group(1))
    for m in re.finditer(r"KNOWN_HTML_SHEETS = new Set\(\[([\s\S]*?)\]\)", text):
        block = m.group(1)
        for c in re.findall(r"'([^']+)'", block):
            codes.add(c)
    return codes


def parse_features(path: Path) -> dict[str, bool]:
    text = path.read_text(encoding="utf-8") if path.exists() else ""
    return {
        "entry_dual_mode": "useWorkpaperEntryDualMode" in text or "EntryDualMode" in text,
        "review_provide": "useWorkpaperReviewProvide" in text,
        "entry_injections": "useWorkpaperEntryInjections" in text,
        "browse_mode_any": "useWorkpaperBrowseMode" in text,
        "import_export": "ImportExport" in text or "importExport" in text,
        "ai_generate": "useAiGenerate" in text or "AiGenerate" in text,
        "get_row_dot": "getRowDot" in text,
    }


def scan_tab_browse(d_dir: Path) -> dict[str, bool]:
    out = {}
    if not d_dir.exists():
        return out
    for f in d_dir.rglob("*.vue"):
        if "Tab" in f.name:
            out[f.name] = "useWorkpaperBrowseMode" in f.read_text(encoding="utf-8")
    return out


def main():
    report: dict = {"cycles": {}, "d4_benchmark": {}, "summary": {}}

    gt_map = {
        "D1": FRONTEND / "GtD1NotesReceivable.vue",
        "D2": FRONTEND / "GtD2AccountsReceivable.vue",
        "D3": FRONTEND / "GtD3PrepaidAccounts.vue",
        "D4": FRONTEND / "GtD4OperatingRevenue.vue",
        "D5": FRONTEND / "GtD5ReceivablesFinancing.vue",
        "D6": FRONTEND / "GtD6ContractAssets.vue",
        "D7": FRONTEND / "GtD7ContractLiabilities.vue",
    }
    label_map = {
        "D1": FRONTEND / "composables" / "d1SheetLabels.ts",
        "D2": FRONTEND / "d2" / "D2TabIndex.vue",
        "D3": FRONTEND / "composables" / "d3SheetLabels.ts",
        "D4": FRONTEND / "composables" / "d4SheetLabels.ts",
        "D5": FRONTEND / "composables" / "d5SheetLabels.ts",
        "D6": FRONTEND / "composables" / "d6SheetLabels.ts",
        "D7": FRONTEND / "composables" / "d7SheetLabels.ts",
    }

    d4_feat = parse_features(gt_map["D4"])
    report["d4_benchmark"] = d4_feat

    for wp in ["D1", "D2", "D3", "D4", "D5", "D6", "D7"]:
        xlsx = read_xlsx_sheets(wp)
        xlsx_codes = {r["code"] for r in xlsx if r.get("code")}
        index_rows = parse_index_rows_ts(label_map[wp])
        index_codes = {r["code"] for r in index_rows if r["applicable"]}
        routes = parse_gt_routes(gt_map[wp])
        feat = parse_features(gt_map[wp])
        d_dir = FRONTEND / f"d{wp[1]}"
        browse = scan_tab_browse(d_dir)

        # xlsx codes not in index
        skip_codes = {"GT_Custom", "directory", "业务模式提示", "修订前程序表", "D7A原", "D8A原", None}
        substantive_xlsx = {c for c in xlsx_codes if c not in skip_codes}
        missing_index = sorted(substantive_xlsx - index_codes - routes)
        missing_route = sorted(index_codes - routes - {"skip"})
        missing_xlsx = sorted(index_codes - substantive_xlsx)

        report["cycles"][wp] = {
            "xlsx_sheet_count": len(xlsx),
            "xlsx_codes": sorted(c for c in xlsx_codes if c),
            "index_codes": sorted(index_codes),
            "html_routes": sorted(routes),
            "features_vs_d4": {k: feat[k] == d4_feat[k] for k in d4_feat},
            "feature_detail": feat,
            "gaps": {
                "xlsx_not_in_frontend": missing_index,
                "index_not_routed": missing_route,
                "index_not_in_xlsx": missing_xlsx,
            },
            "tabs_with_browse_mode": sum(1 for v in browse.values() if v),
            "tabs_total": len(browse),
            "browse_tabs": sorted(k for k, v in browse.items() if v),
            "sheets": xlsx,
            "index_rows": index_rows,
        }

    # parity score
    for wp, data in report["cycles"].items():
        gaps = data["gaps"]
        gap_count = len(gaps["xlsx_not_in_frontend"]) + len(gaps["index_not_routed"]) + len(gaps["index_not_in_xlsx"])
        feat_match = sum(1 for v in data["features_vs_d4"].values() if v)
        data["parity_score"] = {
            "structure_gaps": gap_count,
            "shell_features_matched": f"{feat_match}/{len(d4_feat)}",
        }

    out_path = ROOT / "backend" / "data" / "d_cycle_parity_audit.json"
    out_path.write_text(json.dumps(report, ensure_ascii=False, indent=2), encoding="utf-8")

    # console summary
    print("=" * 72)
    print("D-Cycle Parity Audit (excl D0) — vs D4 benchmark")
    print("=" * 72)
    print(f"D4 shell features: {d4_feat}\n")
    for wp in ["D1", "D2", "D3", "D4", "D5", "D6", "D7"]:
        d = report["cycles"][wp]
        print(f"--- {wp} ---")
        print(f"  xlsx sheets: {d['xlsx_sheet_count']} | index applicable: {len(d['index_codes'])} | html routes: {len(d['html_routes'])}")
        print(f"  browseMode tabs: {d['tabs_with_browse_mode']}/{d['tabs_total']}")
        print(f"  shell parity: {d['parity_score']['shell_features_matched']} | structure gaps: {d['parity_score']['structure_gaps']}")
        g = d["gaps"]
        if g["xlsx_not_in_frontend"]:
            print(f"  [GAP] xlsx no frontend: {g['xlsx_not_in_frontend']}")
        if g["index_not_routed"]:
            print(f"  [GAP] index no route: {g['index_not_routed']}")
        if g["index_not_in_xlsx"]:
            print(f"  [GAP] index no xlsx: {g['index_not_in_xlsx']}")
        diff_feat = [k for k, v in d["features_vs_d4"].items() if not v]
        if diff_feat and wp != "D4":
            print(f"  [GAP] shell diff: {diff_feat}")
        print()

    print(f"Full JSON: {out_path}")


if __name__ == "__main__":
    main()

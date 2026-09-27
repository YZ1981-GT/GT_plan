"""一次性（C-6）：倒出指定册里**兄弟 sheet** 的表头，判断「前端独有字段」的真正归属。

用法：python backend/scripts/analyze/_g_sibling_sheet_probe.py "G3 应收股利.xlsx" "测算及检查表G3-4"
"""
from __future__ import annotations

import sys
from pathlib import Path

from openpyxl import load_workbook
from openpyxl.utils import get_column_letter

TPL = Path(__file__).resolve().parents[2] / "wp_templates" / "G"


def dump(book: str, sheet: str, max_rows: int = 16) -> None:
    wb = load_workbook(TPL / book, data_only=False)
    ws = wb[sheet]
    print(f"\n===== {book} · {sheet} ({ws.max_row}×{ws.max_column}) =====")
    for r in range(1, min(ws.max_row, max_rows) + 1):
        cells = []
        for c in range(1, ws.max_column + 1):
            v = ws.cell(row=r, column=c).value
            if v is None or (isinstance(v, str) and not v.strip()):
                continue
            mark = "ƒ" if isinstance(v, str) and v.startswith("=") else ""
            cells.append(f"{get_column_letter(c)}{mark}={str(v)[:30]}")
        if cells:
            print(f"  R{r:<3} " + " | ".join(cells[:16]) + (" …" if len(cells) > 16 else ""))
    wb.close()


if __name__ == "__main__":
    dump(sys.argv[1], sys.argv[2], int(sys.argv[3]) if len(sys.argv) > 3 else 16)

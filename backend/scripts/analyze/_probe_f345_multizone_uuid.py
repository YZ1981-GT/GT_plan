#!/usr/bin/env python
"""一次性探针（用完即删）：同 sheet 多区的 uuid_col 可用性实测。

spec: f3/f4/f5-sync-coverage-and-first-canary · Task 2

🔴 为什么需要它：同 sheet 有多个受管 table 时，框架层
`phase5_row_table_sheet.spec_to_contract_sheet_payload` 用 **`uuid_col`** 把 spec 与
contract table 一一配对（源码注释：「缺它 → 匹配 0 张 → ProviderCapabilityError」）。
生产先例 D3-4 双区用 J/K、D3-7 双区用 R/S —— **逐区独立空列**。

而三份 spec 的 design 受管区清单把同 sheet 各区的 UUID 列写成**同一列**：
  F3-7 三区全写 S · F4-7 五区全写 L · F4-8 双区全写 S · F4-1 两区全写 M
⇒ 那是不可用的声明。本探针为每张多区 sheet 实测：
  ① 各区的数据行区间（由 footer 的 SUM 区间反推）
  ② 全 sheet 的列占用（哪些列在**任何**行都空）
  ③ 可用作 uuid_col 的候选列数量是否 ≥ 区数
"""
from __future__ import annotations

import argparse
import pathlib
import sys

from openpyxl import load_workbook
from openpyxl.utils import get_column_letter

_BACKEND = pathlib.Path(__file__).resolve().parent.parent.parent
TEMPLATES = {
    "F3": _BACKEND / "wp_templates" / "F" / "F3 应付票据.xlsx",
    "F4": _BACKEND / "wp_templates" / "F" / "F4 应付账款.xlsx",
    "F5": _BACKEND / "wp_templates" / "F" / "F5 营业成本.xlsx",
}

#: 待核的多区 sheet（区数取自 spec，本探针验证其可行性）。
MULTIZONE = {
    "F3": [("应付票据检查表F3-7", 3)],
    "F4": [
        ("未入账检查表F4-7", 5),
        ("应付账款检查表F4-8", 2),
        ("审定表F4-1", 2),
    ],
    "F5": [("营业务成本审定表F5-1", 2)],
}


def probe(ws, zone_count: int) -> None:
    print(f"\n{'=' * 78}")
    print(f"SHEET {ws.title!r}  {ws.max_row}r x {ws.max_column}c "
          f"({get_column_letter(ws.max_column)})  预期 {zone_count} 区")
    print("=" * 78)

    # ① A 列锚点（区标题 / 表头 / footer）
    print("-- A 列锚点（有文本的行）--")
    for r in range(1, ws.max_row + 1):
        v = ws.cell(row=r, column=1).value
        if v is not None and str(v).strip() and not str(v).startswith("="):
            print(f"   R{r:<4} {str(v)[:60]!r}")

    # ② footer 行的 SUM 区间（反推各区数据行）
    print("\n-- 含 SUM 的公式（反推各区数据行区间）--")
    for row in ws.iter_rows():
        for c in row:
            if isinstance(c.value, str) and "SUM(" in c.value.upper():
                print(f"   {c.column_letter}{c.row} = {c.value[:70]}")

    # ③ 列占用（全空列 = uuid_col 候选）
    print("\n-- 列占用 --")
    counts = []
    limit = min(ws.max_column + 8, 60)
    for ci in range(1, limit + 1):
        n = sum(
            1 for r in range(1, ws.max_row + 1) if ws.cell(row=r, column=ci).value is not None
        )
        counts.append((get_column_letter(ci), n))
    print("   " + "  ".join(f"{L}:{n}" for L, n in counts))
    empties = [L for L, n in counts if n == 0]
    print(f"   全空列 ({len(empties)}): {empties}")
    verdict = "✅ 足够" if len(empties) >= zone_count else "🔴 不足"
    print(f"   {verdict} —— 需 {zone_count} 个互不相同的 uuid_col，实得 {len(empties)} 个候选")
    if len(empties) >= zone_count:
        print(f"   建议分配（按区顺序）: {empties[:zone_count]}")

    # ④ 合并区（表头层级判定）
    merged = sorted(str(m) for m in ws.merged_cells.ranges)
    print(f"\n-- 合并区 {len(merged)} 个（前 40）--")
    print("   " + ", ".join(merged[:40]))


def main() -> int:
    ap = argparse.ArgumentParser()
    ap.add_argument("code", choices=sorted(TEMPLATES))
    args = ap.parse_args()
    wb = load_workbook(TEMPLATES[args.code], data_only=False)
    try:
        for name, zones in MULTIZONE[args.code]:
            probe(wb[name], zones)
    finally:
        wb.close()
    return 0


if __name__ == "__main__":
    sys.exit(main())

import openpyxl
wb = openpyxl.load_workbook("backend/wp_templates/M/M5 盈余公积.xlsx", read_only=False, data_only=False)
ws = wb["明细表M5-2"]
# J 列 = column 10
for r in range(7, 17):
    cell = ws.cell(r, 10)
    v = cell.value
    dt = cell.data_type
    vstr = str(v)[:60] if v is not None else "(empty)"
    print(f"J{r}: type={dt} value={vstr}")
# 同时检查 H, I, K, L, M 列（都是我声明为 editable 的）
print()
for col_idx, col_name in [(8, "H"), (9, "I"), (11, "K"), (12, "L"), (13, "M")]:
    for r in [12, 13, 14, 15, 16]:
        cell = ws.cell(r, col_idx)
        v = cell.value
        vstr = str(v)[:40] if v is not None else "(empty)"
        is_f = isinstance(v, str) and v.startswith("=")
        print(f"{col_name}{r}: {'F' if is_f else 'V'} {vstr}")
wb.close()

"""Offline field/column reconciliation for the S-checklist templates.

The S checklist components are HTML views of several legacy workbooks.  This
checker deliberately chooses the reader from the workbook suffix:
``.xlsx`` is read with openpyxl and the legacy S17 ``.xls`` is read with xlrd.
It never opens an ``.xls`` file with openpyxl.

The output is a deterministic JSON report.  Differences are evidence for the
component/template contract review; this tool does not modify templates or the
database.

Usage from the repository root::

    python backend/scripts/analyze/check_s_checklist_template_alignment.py \
      --output .agents/tasks/s-cycle-p0-impl-2026-10-25/s_checklist_template_alignment.json
"""
from __future__ import annotations

import argparse
import json
import re
from pathlib import Path
from typing import Any, Iterable

ROOT = Path(__file__).resolve().parents[3]
TEMPLATE_ROOT = ROOT / "backend" / "wp_templates" / "S"


def _norm(value: Any) -> str:
    """Normalize a cell for header comparison without changing the raw report."""
    return re.sub(r"\s+", "", str(value or "")).strip()


def _relative(path: Path) -> str:
    try:
        return path.resolve().relative_to(ROOT.resolve()).as_posix()
    except ValueError:
        return path.resolve().as_posix()


def _non_empty(values: Iterable[Any]) -> list[str]:
    return [str(value).strip() for value in values if value not in (None, "")]


def _read_xlsx(path: Path) -> dict[str, Any]:
    """Read an xlsx workbook through openpyxl and return plain values only."""
    import openpyxl

    workbook = openpyxl.load_workbook(path, read_only=True, data_only=True)
    sheets: dict[str, dict[str, Any]] = {}
    try:
        for worksheet in workbook.worksheets:
            rows = [list(row) for row in worksheet.iter_rows(values_only=True)]
            sheets[worksheet.title] = {
                "rows": rows,
                "rows_count": worksheet.max_row,
                "columns_count": worksheet.max_column,
            }
    finally:
        workbook.close()
    return {"reader": "openpyxl", "format": "xlsx", "sheets": sheets}


def _read_xls(path: Path) -> dict[str, Any]:
    """Read a legacy xls workbook through xlrd, never through openpyxl."""
    import xlrd

    workbook = xlrd.open_workbook(str(path), formatting_info=False)
    sheets: dict[str, dict[str, Any]] = {}
    for worksheet in workbook.sheets():
        sheets[worksheet.name] = {
            "rows": [worksheet.row_values(row_index) for row_index in range(worksheet.nrows)],
            "rows_count": worksheet.nrows,
            "columns_count": worksheet.ncols,
        }
    return {
        "reader": "xlrd",
        "reader_version": str(getattr(xlrd, "__version__", "unknown")),
        "format": "xls",
        "sheets": sheets,
    }


# The labels below are the visible columns/fields rendered by the current
# components.  Aliases capture the small wording differences that are valid
# across the HTML view and the workbook (for example 本期/本年度).
_CASES: tuple[dict[str, Any], ...] = (
    {
        "component_type": "s1-regulation-record",
        "workbook": "S1 对被审计单位违反法规行为的考虑2019.xlsx",
        "sheets": ("S1-1对法律法规的考虑记录表",),
        "comparison": "freeform_template_vs_grid_component",
        "ui_columns": (
            ("index", "序号", ("序号",)),
            ("category", "法规类别", ("法规类别",)),
            ("lawName", "法律法规名称", ("法律法规名称", "法规名称")),
            ("applicable", "是否适用", ("是否适用",)),
            ("impactAssessment", "对财务报表的影响", ("对财务报表的影响",)),
            ("consequence", "不合规后果", ("不合规后果",)),
            ("auditProcedure", "已执行审计程序", ("已执行审计程序",)),
            ("conclusion", "结论", ("结论",)),
            ("indexRef", "索引号", ("索引号",)),
        ),
        "header_groups": (),
    },
    {
        "component_type": "s8-adjudication",
        "workbook": "S8 租赁.xlsx",
        "sheets": ("审定表S8-1",),
        "comparison": "freeform_template_vs_grid_component",
        "ui_columns": (
            ("accountName", "科目名称", ("科目名称",)),
            ("accountCode", "科目编码", ("科目编码",)),
            ("unadjustedAmount", "未审数", ("未审数",)),
            ("adjustmentAmount", "审计调整", ("审计调整",)),
            ("auditedAmount", "审定数", ("审定数",)),
            ("ajeRef", "调整分录", ("调整分录",)),
            ("indexRef", "索引号", ("索引号",)),
        ),
        "header_groups": (),
    },
    {
        "component_type": "s9-adjudication",
        "workbook": "S9 对电子商务的考虑.xlsx",
        "sheets": ("审定表S9-1",),
        "comparison": "freeform_template_vs_grid_component",
        "ui_columns": (
            ("accountName", "科目名称", ("科目名称",)),
            ("accountCode", "科目编码", ("科目编码",)),
            ("unadjustedAmount", "未审数", ("未审数",)),
            ("adjustmentAmount", "审计调整", ("审计调整",)),
            ("auditedAmount", "审定数", ("审定数",)),
            ("ajeRef", "调整分录", ("调整分录",)),
            ("indexRef", "索引号", ("索引号",)),
        ),
        "header_groups": (),
    },
    {
        "component_type": "s9-internal-control",
        "workbook": "S9 对电子商务的考虑.xlsx",
        "sheets": ("内部控制调查表S9-2",),
        "comparison": "table_header",
        "ui_columns": (
            ("controlItem", "控制活动/调查事项", ("控制活动/调查事项", "调查问题")),
            ("rating", "评价", ("是", "弱", "否", "不适用")),
            ("evidence", "备注/审计证据", ("备注/审计证据", "备注")),
            ("indexRef", "索引号", ("索引号",)),
        ),
        "header_groups": (("调查问题", "答案"), ("是", "弱", "否", "不适用")),
    },
    {
        "component_type": "s9-legal-regulations",
        "workbook": "S9 对电子商务的考虑.xlsx",
        "sheets": ("电子商务法律法规",),
        "comparison": "readonly_reference",
        "ui_columns": (),
        "header_groups": (),
    },
    {
        "component_type": "s9-program-description",
        "workbook": "S9 对电子商务的考虑.xlsx",
        "sheets": ("程序说明",),
        "comparison": "readonly_reference",
        "ui_columns": (),
        "header_groups": (),
    },
    {
        "component_type": "s10-adjudication",
        "workbook": "S10 对环境事项的考虑.xlsx",
        "sheets": ("审定表S10-1",),
        "comparison": "freeform_template_vs_grid_component",
        "ui_columns": (
            ("accountName", "科目名称", ("科目名称",)),
            ("accountCode", "科目编码", ("科目编码",)),
            ("unadjustedAmount", "未审数", ("未审数",)),
            ("adjustmentAmount", "审计调整", ("审计调整",)),
            ("auditedAmount", "审定数", ("审定数",)),
            ("ajeRef", "调整分录", ("调整分录",)),
            ("indexRef", "索引号", ("索引号",)),
        ),
        "header_groups": (),
    },
    {
        "component_type": "s10-internal-control",
        "workbook": "S10 对环境事项的考虑.xlsx",
        "sheets": ("附、内控调查表S10-2",),
        "comparison": "table_header",
        "ui_columns": (
            ("controlItem", "控制活动/调查事项", ("控制活动/调查事项", "调查问题")),
            ("rating", "评价", ("是", "弱", "否", "不适用")),
            ("evidence", "备注/审计证据", ("备注/审计证据", "备注")),
            ("indexRef", "索引号", ("索引号",)),
        ),
        "header_groups": (("调查问题", "答案"), ("是", "弱", "否", "不适用")),
    },
    {
        "component_type": "s10-environmental-regulations",
        "workbook": "S10 对环境事项的考虑.xlsx",
        "sheets": ("环境法规",),
        "comparison": "readonly_reference",
        "ui_columns": (),
        "header_groups": (),
    },
    {
        "component_type": "s17-nonrecurring-detail",
        "workbook": "S17 非经常性损益.xls",
        "sheets": (
            "审定表S17-1",
            "审核表（合并）S17-2",
            "审核表（个别报表）S17-3",
        ),
        "comparison": "table_header",
        "ui_columns": {
            "审定表S17-1": (
                ("itemName", "非经常性损益项目", ("非经常性损益项目", "项目", "项  目")),
                ("currentAmount", "本期金额", ("本期金额", "本年度")),
                ("priorAmount", "上期金额", ("上期金额", "上年度")),
                ("taxImpact", "所得税影响", ("所得税影响",)),
                ("netAmount", "税后净额", ("税后净额",)),
                ("indexRef", "索引号", ("索引号",)),
            ),
            "审核表（合并）S17-2": (
                ("index", "序号", ("序号",)),
                ("itemName", "非经常性损益项目", ("非经常性损益项目", "项目", "项  目")),
                ("amount", "金额", ("金额", "合并金额")),
                ("source", "来源科目", ("来源科目",)),
                ("basis", "认定依据", ("认定依据",)),
                ("crossRef", "勾稽", ("勾稽",)),
                ("auditResult", "审核结果", ("审核结果",)),
            ),
            "审核表（个别报表）S17-3": (
                ("index", "序号", ("序号",)),
                ("itemName", "非经常性损益项目", ("非经常性损益项目", "项目", "项  目")),
                ("amount", "金额", ("金额",)),
                ("source", "来源科目", ("来源科目",)),
                ("basis", "认定依据", ("认定依据", "计算说明")),
                ("auditResult", "审核结果", ("审核结果",)),
            ),
        },
        "header_groups": {
            "审定表S17-1": (("项目", "本年度"), ("未审数", "审计调整", "审定数")),
            "审核表（合并）S17-2": (("项目", "母公司", "合计", "合并抵销", "合并金额"),),
            "审核表（个别报表）S17-3": (("项目", "金额", "计算说明", "索引号"),),
        },
    },
    {
        "component_type": "s17-reconciliation",
        "workbook": "S17 非经常性损益.xls",
        "sheets": ("与营业外收支等核对表S17-11",),
        "comparison": "table_header",
        "ui_columns": (
            ("itemName", "项目名称", ("项目名称", "项目")),
            ("plAmount", "利润表金额", ("利润表金额",)),
            ("nrAmount", "非经常性损益金额", ("非经常性损益金额", "非经常性损益")),
            ("recurringAmount", "经常性损益金额", ("经常性损益金额",)),
            ("difference", "差异", ("差异",)),
            ("note", "说明", ("说明", "差异原因")),
            ("crossRef", "勾稽", ("勾稽",)),
        ),
        "header_groups": (("项目", "金额", "差异原因"), ("非经常性损益", "其他收益", "投资收益", "资产处置损益", "营业外收支", "差异")),
    },
    {
        "component_type": "s17-tax-impact",
        "workbook": "S17 非经常性损益.xls",
        "sheets": ("非经常性损益项目的所得税影响S17-21",),
        "comparison": "table_header",
        "ui_columns": (
            ("itemName", "非经常性损益项目", ("非经常性损益项目", "项目")),
            ("grossAmount", "税前金额", ("税前金额", "金额")),
            ("deductible", "可抵扣", ("可抵扣",)),
            ("effectiveRate", "适用税率", ("适用税率", "税率")),
            ("taxImpact", "所得税影响", ("所得税影响", "所得税费用")),
            ("netAmount", "税后净额", ("税后净额", "扣非后")),
            ("auditNote", "审核说明", ("审核说明", "说明")),
        ),
        "header_groups": (("项目", "金额", "其中：非经常性损益", "扣非后", "说明"),),
    },
)


def _header_groups_for(case: dict[str, Any], sheet_name: str) -> tuple[tuple[str, ...], ...]:
    groups = case["header_groups"]
    if isinstance(groups, dict):
        return tuple(tuple(group) for group in groups.get(sheet_name, ()))
    return tuple(tuple(group) for group in groups)


def _find_header_rows(rows: list[list[Any]], groups: tuple[tuple[str, ...], ...]) -> list[int]:
    found: list[int] = []
    for group in groups:
        wanted = tuple(_norm(value) for value in group)
        for row_index, row in enumerate(rows, 1):
            values = {_norm(value) for value in row if value not in (None, "")}
            if all(token in values for token in wanted):
                found.append(row_index)
                break
    return sorted(set(found))


def _columns_from_rows(rows: list[list[Any]], header_rows: list[int]) -> list[str]:
    columns: list[str] = []
    seen: set[str] = set()
    for row_number in header_rows:
        for raw in _non_empty(rows[row_number - 1]):
            normalized = _norm(raw)
            if normalized and normalized not in seen:
                seen.add(normalized)
                columns.append(raw)
    return columns


def _ui_columns_for(case: dict[str, Any], sheet_name: str) -> tuple[tuple[str, str, tuple[str, ...]], ...]:
    columns = case["ui_columns"]
    if isinstance(columns, dict):
        return tuple(columns.get(sheet_name, ()))
    return tuple(columns)


def _record_for_sheet(
    case: dict[str, Any],
    workbook: dict[str, Any],
    sheet_name: str,
) -> dict[str, Any]:
    sheet = workbook["sheets"].get(sheet_name)
    if sheet is None:
        return {
            "sheet_name": sheet_name,
            "status": "missing_sheet",
            "reader": workbook["reader"],
            "format": workbook["format"],
            "differences": [{"kind": "missing_template_sheet", "sheet": sheet_name}],
        }

    header_rows = _find_header_rows(
        sheet["rows"],
        _header_groups_for(case, sheet_name),
    )
    template_columns = _columns_from_rows(sheet["rows"], header_rows)
    actual = {_norm(value): value for value in template_columns}
    ui_column_specs = _ui_columns_for(case, sheet_name)
    ui_columns = [
        {"field": field, "label": label, "aliases": list(aliases)}
        for field, label, aliases in ui_column_specs
    ]

    matched_fields: list[dict[str, str]] = []
    consumed: set[str] = set()
    differences: list[dict[str, Any]] = []
    for field, label, aliases in ui_column_specs:
        matching = [
            actual[_norm(alias)]
            for alias in aliases
            if _norm(alias) in actual
        ]
        if matching:
            matched_fields.append({"field": field, "template_label": matching[0]})
            consumed.update(_norm(value) for value in matching)
        else:
            differences.append({
                "kind": "ui_field_missing_in_template_header",
                "field": field,
                "ui_label": label,
                "accepted_template_labels": list(aliases),
            })

    for normalized, raw in actual.items():
        if normalized not in consumed:
            differences.append({
                "kind": "template_column_not_consumed_by_ui",
                "template_label": raw,
            })

    return {
        "sheet_name": sheet_name,
        "status": "ok",
        "reader": workbook["reader"],
        "format": workbook["format"],
        "dimensions": {
            "rows": sheet["rows_count"],
            "columns": sheet["columns_count"],
        },
        "header_rows": header_rows,
        "template_columns": template_columns,
        "ui_columns": ui_columns,
        "matched_fields": matched_fields,
        "differences": differences,
        "column_count_delta": len(ui_columns) - len(template_columns),
        "comparison": case["comparison"],
    }


def build_report(template_root: Path = TEMPLATE_ROOT) -> dict[str, Any]:
    """Build a deterministic report from the five real S templates."""
    cache: dict[str, dict[str, Any]] = {}
    entries: list[dict[str, Any]] = []
    for case in _CASES:
        workbook_name = case["workbook"]
        path = template_root / workbook_name
        if not path.exists():
            entries.append({
                "component_type": case["component_type"],
                "workbook": _relative(path),
                "status": "missing_workbook",
                "differences": [{"kind": "missing_template_file", "path": _relative(path)}],
            })
            continue

        suffix = path.suffix.lower()
        if suffix == ".xlsx":
            reader = "openpyxl"
            workbook_format = "xlsx"
        elif suffix == ".xls":
            reader = "xlrd"
            workbook_format = "xls"
        else:  # pragma: no cover - guarded by the checked-in case table
            raise ValueError(f"Unsupported S template format: {path}")
        if workbook_name not in cache:
            cache[workbook_name] = _read_xlsx(path) if reader == "openpyxl" else _read_xls(path)
        workbook = cache[workbook_name]
        sheet_entries = [
            _record_for_sheet(case, workbook, sheet_name)
            for sheet_name in case["sheets"]
        ]
        entries.append({
            "component_type": case["component_type"],
            "workbook": _relative(path),
            "status": "ok",
            "format": workbook_format,
            "reader": reader,
            "sheets": sheet_entries,
        })

    return {
        "report_version": 1,
        "purpose": "D4.4 S-checklist component/template field and column reconciliation",
        "read_only": True,
        "format_routing": {".xlsx": "openpyxl", ".xls": "xlrd"},
        "component_count": len(entries),
        "component_types": [entry["component_type"] for entry in entries],
        "entries": entries,
    }


def main(argv: list[str] | None = None) -> int:
    parser = argparse.ArgumentParser(description="离线对账 S-checklist 组件字段与真实模板列")
    parser.add_argument("--template-root", type=Path, default=TEMPLATE_ROOT)
    parser.add_argument("--output", type=Path)
    args = parser.parse_args(argv)

    report = build_report(args.template_root)
    encoded = json.dumps(report, ensure_ascii=False, indent=2, sort_keys=True) + "\n"
    if args.output:
        args.output.parent.mkdir(parents=True, exist_ok=True)
        args.output.write_text(encoded, encoding="utf-8")
        print(f"[ALIGNMENT] report={args.output}")
    else:
        print(encoded, end="")
    print(
        f"[ALIGNMENT] components={report['component_count']} "
        f"formats={report['format_routing']}"
    )
    return 0


if __name__ == "__main__":
    raise SystemExit(main())

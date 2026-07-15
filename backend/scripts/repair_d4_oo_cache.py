"""Repair all mismatched D4 OnlyOffice caches for the active project."""
from __future__ import annotations

import shutil
import uuid
from pathlib import Path

from openpyxl import load_workbook

from app.routers.wp_onlyoffice_router import (
    _extract_sheet_code,
    _oo_cache_mismatches_template,
)
from app.services import wp_template_finder as f
from app.services.wp_template_finder import find_template_file_any

f._index_cache = None
pid = uuid.UUID("0ec33ac9-3de5-4e65-b3bf-f9dccd7b2a49")
base = Path("storage/projects") / str(pid) / "workpapers" / "onlyoffice"

for p in sorted(base.glob("D4*.xlsx")):
    if p.name.endswith("__whole.xlsx"):
        continue
    stem = p.stem
    tpl = find_template_file_any(stem)
    if not tpl:
        print("skip no tpl", stem)
        continue
    need = _oo_cache_mismatches_template(p, wp_code=stem, visible_sheet=None)
    if not need and stem != "D4":
        names = load_workbook(p, read_only=True).sheetnames
        codes = {_extract_sheet_code(n) for n in names}
        if stem not in codes:
            need = True
    if need:
        shutil.copy2(tpl, p)
        print("fixed", stem, "<-", tpl.name, p.stat().st_size)
    else:
        print("keep", stem, p.stat().st_size)

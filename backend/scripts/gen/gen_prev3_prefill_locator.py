"""生成三参 PREV 预填逐格定位表 ``backend/data/prev3_prefill_locator.json``。

spec: formula-fetch-followup-defect-closure R2.1 / R2.2

映射表 ``prefill_formula_mapping.json`` 里每个三参 ``PREV('底稿','表','字段')`` 格（全部块，不只第一块）→ 系统模板
（权威目录，不经事务所覆盖层）里的确切单元格与动作。扫描规则（与运行时核对同一套读法，见
``app.services.prev3_prefill_locator``）：

* 行：审定表先取「…原值」分节的第一个小计 / 合计行（三参取数定义 = 「未审数」那组科目，是原值口径），没有原值分节
  取「差异数」行公式里与「试算平衡表数」相减的那一行，再没有取第一个「…合计」行；其余表取第一个「…合计」行。
* 列：列头路径含「上年 / 上期 / 期初 / 年初」且不含比较类字样（比较 / 变动 / 差异 / 结构比 / …）的列，再按字段收窄
  （未审 / 账面价值 / 减值 / 审定 / 余额 / 发生额），多列时取列头有「审定数」整词的那一列；剩一列才算定位成功。
* 动作：目标格是模板公式 ⇒ ``verify``，否则 ``write``；``MANUAL`` 登记的格以人工裁定为准（单元格或 ``none``，写原因）。
  自动规则定不了、又没有人工裁定的格 ⇒ 生成失败并列出，不猜。

输出确定（无时间戳、按键排序，模板文件记 sha256 —— 模板变了 ``--check`` 就红）。用法（仓库根执行）：
    python backend/scripts/gen/gen_prev3_prefill_locator.py            # 写入（内容未变则不改文件）
    python backend/scripts/gen/gen_prev3_prefill_locator.py --check    # 漂移 exit 2 并点名
    python backend/scripts/gen/gen_prev3_prefill_locator.py --report X # 另写一份逐格审阅清单（不入库）
"""
from __future__ import annotations

import argparse
import hashlib
import json
import re
import sys
import warnings
from pathlib import Path
from typing import Any

BACKEND = Path(__file__).resolve().parents[2]
sys.path.insert(0, str(BACKEND))

from app.services.prev3_prefill_locator import (  # noqa: E402
    COMPARISON_TOKENS,
    LOCATOR_FILE,
    PREV3_RE,
    PRIOR_TOKENS,
    header_path,
    is_formula_value,
    merged_values,
    raw_value,
    row_label,
    target_signature,
)

MAPPING = BACKEND / "data" / "prefill_formula_mapping.json"
_SECTION_RE = re.compile(r"^([一二三四五六七八九十]+、|（[一二三四五六七八九十]+）)")
_ROW_REF_RE = re.compile(r"\$?[A-Z]{1,3}\$?([0-9]+)")
_MINUS_RE = re.compile(r"^=\s*\$?[A-Z]{1,3}\$?([0-9]+)\s*-\s*\$?[A-Z]{1,3}\$?([0-9]+)\s*$")
#: 科目原值口径的分节（三参取数定义 = 「未审数」那组科目 = 原值；坏账 / 减值 / 未实现融资收益另列）
_GROSS_SECTIONS = ("原值", "账面余额", "原始金额", "余额", "公允价值")
#: 列头里的「审定数」整词（区别于「披露审定数」「最终审定数」）
_EXACT_AUDITED = ("审定数", "上期审定数", "期初审定数", "上年审定数", "上年同期审定数", "审定余额", "期初审定余额")
_MAX_ROW, _MAX_COL = 250, 60

#: 人工裁定（自动规则定不了或定错的格）：``(底稿, 映射表 sheet, cell_ref) → {"cell": A1, "reason": …}``
#: 或 ``{"action": "none", "reason": …}``。每条都必须写原因；失效条目（映射表里已没有该格）生成时报错。
#: 2026-10-24 逐格看模板裁定（审阅清单见 spec design §10.9）。
_NO_PRIOR_TABLE = "本表没有上年 / 期初金额栏"
MANUAL: dict[tuple[str, str, str], dict[str, str]] = {
    # ── 审定表：结构特殊，自动规则取不到科目原值合计 ──
    ("F2", "存货审定表F2-1", "上年审定数"): {
        "cell": "B35", "reason": "存货审定表按「未审数 / 账项调整 / 审定数」分段，没有小计行；取「一、原值」段的审定数（原值）期初数"},
    ("G7", "长期股权投资审定表G7-1", "上年审定数"): {
        "cell": "C22", "reason": "长期股权投资审定表按「未审数 / 账项调整 / 重分类调整 / 审定数」分段；取「一、原值」段的审定数1 期初数"},
    ("J1", "审定表J1-1", "上年审定数"): {
        "action": "none", "reason": "应付职工薪酬审定表分短期薪酬 / 离职后福利 / 辞退福利三张子表，没有单一合计行"
                                    "（模板差异数 = 三表合计之和 − 试算平衡表数）"},
    # ── 明细表 / 附注：多列候选，按格名与取数口径裁定 ──
    ("D2", "坏账准备明细表D2-3", "上年坏账准备余额"): {
        "cell": "E22", "reason": "上年坏账准备期末余额 = 本年期初审定数（期初数 / 期初审定数列）"},
    ("D4", "附注披露信息（上市公司）", "上年主营业务收入"): {"cell": "D10", "reason": "上期发生额 / 收入列的「主营业务」行"},
    ("D4", "附注披露信息（上市公司）", "上年营业成本合计"): {"cell": "E12", "reason": "上期发生额 / 成本列的合计行"},
    ("D4", "附注披露信息（国企）", "上年主营业务收入"): {"cell": "D9", "reason": "上期发生额 / 收入列的「主营业务」行"},
    ("D4", "附注披露信息（国企）", "上年营业成本合计"): {"cell": "E11", "reason": "上期发生额 / 成本列的合计行"},
    ("D5", "附注披露信息（上市公司）", "上年账面价值"): {
        "cell": "C13", "reason": "第一张表没有合计行，账面价值 = 「期末公允价值」行；第一个合计行（第 33 行）是已质押票据表"},
    ("H1", "明细表H1-2", "固定资产净值_上年审定"): {"cell": "AV33", "reason": "格名是净值：取「期初净值 / 审定净值」列合计"},
    ("H3", "明细表（成本模式）H3-2", "上年明细审定"): {
        "cell": "L28", "reason": "成本模式下科目余额是原值（累计折旧 / 减值准备另列）：取原值 / 审定数 / 期初余额列合计"},
    ("H3", "明细表（公允价值模式）H3-2", "上年明细审定"): {
        "cell": "X27", "reason": "公允价值模式下科目余额 = 成本 + 公允价值变动：取「期初净值」列合计"},
    ("H8", "明细表H8-2", "上年明细审定"): {
        "cell": "S32", "reason": "使用权资产科目余额是原值（累计折旧 / 减值准备另列）：取原值 / 审定数 / 期初数列合计"},
    ("I3", "明细表I3-2", "上年商誉_审定"): {
        "cell": "M23", "reason": "商誉科目余额是原值（减值准备另列，另一格「上年商誉减值_审定」取减值列）：取商誉原值 / 审定数 / 期初数列合计"},
    ("M4", "明细表M4-2", "资本溢价_审定期初"): {"cell": "L10", "reason": "审定数 / 期初数列的「一、资本（股本）溢价」行"},
    ("M4", "明细表M4-2", "其他资本公积_审定期初"): {"cell": "L15", "reason": "审定数 / 期初数列的「二、其他资本公积」行"},
    ("M6", "明细表M6-2", "上年年末余额_审定数"): {"cell": "E10", "reason": "本期数 / 审定数列的「一、上年年末余额」行（上年语义在行标签）"},
    ("M6", "明细表M6-2", "本年年初余额_审定数"): {
        "action": "none", "reason": "本年年初余额 = 上年年末余额 + 会计政策变更 + 前期差错更正，三参取的是上年年末余额，"
                                    "不对应本格；上年年末余额已由「上年年末余额_审定数」格核对"},
    ("G7", "明细表G7-2", "上年审定数"): {
        "action": "none", "reason": "明细表按成本法 / 权益法两张子表叠放，审定数「期初余额」列头只在投资比例列，金额列没有期初字样；"
                                    "上年审定数已由审定表 G7-1 核对"},
    # ── 同一目标格、取数参数不同：只保留一格 ──
    ("F5", "营业务成本审定表F5-1", "上年发生额"): {
        "action": "none", "reason": "与「上年审定数」落同一格（上期数 / 上期审定数合计）而取数参数不同，以「上年审定数」为准"},
    ("K8", "实质性分析K8-4", "本年期初"): {
        "action": "none", "reason": "与「上年发生额」落同一格（上年同期合计）而取数参数不同；费用类没有「本年期初」口径"},
    # ── 模板没有对应格 ──
    ("D0", "跟函函证过程控制D0-3", "上年审定数"): {"action": "none", "reason": "跟函过程记录是文字备忘录，没有金额栏"},
    ("D2", "应收账款分析表D2-5", "上年审定数"): {"action": "none", "reason": f"{_NO_PRIOR_TABLE}（借贷方发生额核对）"},
    ("D3", "预收账款分析表D3-4", "上年审定数"): {"action": "none", "reason": f"{_NO_PRIOR_TABLE}（借方发生额分析）"},
    ("D7", "合同负债分析表D7-4", "上年审定数"): {"action": "none", "reason": f"{_NO_PRIOR_TABLE}（借方发生额分析）"},
    ("D5", "应收款项融资公允价值测算表D5-4", "上年公允价值合计"): {"action": "none", "reason": f"{_NO_PRIOR_TABLE}（期末公允价值测算）"},
    ("E1", "货币资金分析表E1-14", "分析表上年审定数"): {
        "action": "none", "reason": "构成分析表列头是固定日期（2025 / 2024 / 2023-12-31），没有「上年 / 期初」字样"},
    ("F2", "存货总体分析表F2-18", "上年审定数（分析程序）"): {
        "action": "none", "reason": "构成分析表列头是年份占位（20X5 / 20X4 / 20X3），没有「上年 / 期初」字样；金额由明细汇总表公式引用"},
    ("H1", "折旧分配分析表H1-13", "上年折旧分配审定"): {"action": "none", "reason": f"{_NO_PRIOR_TABLE}（本期折旧按费用项目分配）"},
    ("H1", "减值测算表H1-14", "上年减值测算审定"): {"action": "none", "reason": f"{_NO_PRIOR_TABLE}（期末减值测算）"},
    ("J1", "调整分录汇总表J1-3", "上年审定数"): {"action": "none", "reason": "调整分录汇总表是分录清单，没有金额合计栏"},
    ("J1", "月度分析表J1-4", "上年同期期末"): {"action": "none", "reason": "月度分析表按月列示本期 / 上期计提数，没有期末余额栏"},
    ("J1", "月度分析表J1-4", "上年同期费用率"): {"action": "none", "reason": "月度分析表没有费用率栏"},
    ("J3", "股份支付情况表J3-1", "上年审定数"): {"action": "none", "reason": "股份支付情况表是文字与清单，没有金额合计行（股份支付底稿没有审定表）"},
}


# ── 模板 ─────────────────────────────────────────────────────────────────────────

_books: dict[str, list[tuple[Path, Any]]] = {}


def _workbooks(wp_code: str) -> list[tuple[Path, Any]]:
    """该底稿的系统模板（主文件在前，同组其余 xlsx 在后；与初始化合并 sheet 的先后一致）。"""
    if wp_code not in _books:
        from openpyxl import load_workbook

        from app.services.wp_template_finder import (
            find_all_template_files_unresolved,
            find_template_file_any_unresolved,
        )

        paths: list[Path] = []
        main = find_template_file_any_unresolved(wp_code)
        if main:
            paths.append(main)
        for p in find_all_template_files_unresolved(wp_code) or []:
            if p not in paths:
                paths.append(p)
        with warnings.catch_warnings():
            warnings.simplefilter("ignore")
            _books[wp_code] = [(p, load_workbook(str(p))) for p in paths if p.suffix.lower() == ".xlsx"]
    return _books[wp_code]


def _find_sheet(wp_code: str, sheet: str) -> tuple[Path, Any] | None:
    for path, wb in _workbooks(wp_code):
        for ws in wb.worksheets:
            if ws.title.strip() == sheet:
                return path, ws
    return None


def _rel(path: Path) -> str:
    return path.resolve().relative_to(BACKEND).as_posix()


def _sha256(path: Path) -> str:
    return hashlib.sha256(path.read_bytes()).hexdigest()


# ── 自动定位 ──────────────────────────────────────────────────────────────────────


def _is_total(label: str) -> bool:
    text = label.rstrip("：:")
    return text.endswith("合计") or text == "总计"


def _is_subtotal(label: str) -> bool:
    return label.rstrip("：:").endswith("小计") or _is_total(label)


def _is_section(ws: Any, row: int, label: str) -> bool:
    """分节标题行：「一、…」「（一）…」或以冒号结尾，且本行没有取数公式 / 数值。"""
    if not (_SECTION_RE.match(label) or label.endswith(("：", ":"))):
        return False
    return not any(isinstance(v, (int, float)) or is_formula_value(v)
                   for v in (raw_value(ws, row, c) for c in range(2, min(ws.max_column, _MAX_COL) + 1)))


def _first_leaf_section(ws: Any, labels: dict[int, str]) -> int | None:
    """第一个分节标题；紧跟着又是分节标题（「一、其他应收款」+「（一）其他应收款原值」）时取后者。"""
    heads = [r for r, lab in labels.items() if lab and _is_section(ws, r, lab)]
    if not heads:
        return None
    return heads[1] if len(heads) > 1 and heads[1] == heads[0] + 1 else heads[0]


def _diff_target(ws: Any, labels: dict[int, str], row: int) -> int | None:
    """「差异数」行公式里与「试算平衡表数」相减的合计行（行标签须是「…合计 / 小计」或公式标签，否则不认）。"""
    for c in range(2, min(ws.max_column, _MAX_COL) + 1):
        value = raw_value(ws, row, c)
        if not is_formula_value(value):
            continue
        for ref in (int(x) for x in _ROW_REF_RE.findall(value)):
            label = labels.get(ref, "")
            if ref < row and (_is_subtotal(label) or label == ""):
                return ref
        return None
    return None


def _before_deduction(ws: Any, labels: dict[int, str], row: int) -> int:
    """合计行是「小计 − 减：…」⇒ 退回扣减前的小计（列报扣减不在科目余额里）。"""
    for c in range(2, min(ws.max_column, _MAX_COL) + 1):
        m = _MINUS_RE.match(str(raw_value(ws, row, c) or ""))
        if m and labels.get(int(m.group(2)), "").startswith("减"):
            return int(m.group(1))
    return row


def account_row(ws: Any, merged: dict, audit_sheet: bool) -> tuple[int | None, str]:
    """科目合计所在行 →（行号, 规则名）。审定表：原值 / 账面余额分节的第一个小计 → 差异数核对行（扣减前）→ 首个合计行。"""
    rows = range(4, min(ws.max_row, _MAX_ROW) + 1)
    labels = {r: row_label(ws, r, merged) for r in rows}
    if audit_sheet:
        head = _first_leaf_section(ws, labels)
        core = _SECTION_RE.sub("", labels.get(head, "")).rstrip("：:") if head else ""
        if head and core.endswith(_GROSS_SECTIONS):
            for r in rows:
                if r <= head:
                    continue
                if _is_section(ws, r, labels[r]):
                    break
                if _is_subtotal(labels[r]):
                    return r, "原值小计"
        for r in rows:
            if labels[r].startswith("差异"):
                target = _diff_target(ws, labels, r)
                if target is not None:
                    return _before_deduction(ws, labels, target), "差异数核对行"
                break
    for r in rows:
        if _is_total(labels[r]):
            return r, "首个合计行"
    return None, ""


def _body(header: tuple[str, ...]) -> list[str]:
    """列头里去掉表名（以「表」结尾的标题行）后的部分 —— 语义判断只看它。"""
    return [p for p in header if not p.endswith("表")]


def is_prior_header(header: tuple[str, ...]) -> bool:
    """列头属上年 / 期初语义且不是比较类列（守卫与生成器同口径）。"""
    body = _body(header)
    return any(t in p for p in body for t in PRIOR_TOKENS) and not any(t in p for p in body for t in COMPARISON_TOKENS)


def _has(header: tuple[str, ...], *tokens: str) -> bool:
    return any(t in p for p in _body(header) for t in tokens)


def prior_columns(ws: Any, row: int, merged: dict) -> dict[int, tuple[str, ...]]:
    out: dict[int, tuple[str, ...]] = {}
    for col in range(2, min(ws.max_column, _MAX_COL) + 1):
        header = header_path(ws, col, row, merged)
        if header and is_prior_header(header):
            out[col] = header
    return out


def pick_column(cols: dict[int, tuple[str, ...]], field: str) -> list[int]:
    """按三参第三参（字段）收窄候选列；多列时取列头有「审定数」整词的那列；收窄后为空而只有一个上年列就用它。"""
    pool = sorted(cols)
    if "未审" in field:
        keep = [c for c in pool if _has(cols[c], "未审")]
    elif "账面价值" in field:
        keep = [c for c in pool if _has(cols[c], "账面价值")]
    elif "减值" in field:
        keep = [c for c in pool if _has(cols[c], "减值") and _has(cols[c], "审定")]
    elif "审定" in field:
        keep = [c for c in pool if _has(cols[c], "审定")]
    elif "余额" in field:
        keep = [c for c in pool if _has(cols[c], "余额")]
    elif "发生额" in field or field.endswith("本期"):
        keep = [c for c in pool if _has(cols[c], "发生额")]
    else:
        keep = [c for c in pool if _has(cols[c], "审定")]
    if len(keep) > 1:
        exact = [c for c in keep if any(p in _EXACT_AUDITED for p in _body(cols[c]))]
        keep = exact or keep
    if not keep and len(pool) == 1:
        keep = pool
    return keep


def auto_locate(ws: Any, field: str, audit_sheet: bool) -> tuple[str | None, str, str]:
    """→（A1, 规则, 定不了的原因）。"""
    from openpyxl.utils import get_column_letter

    merged = merged_values(ws)
    row, rule = account_row(ws, merged, audit_sheet)
    if row is None:
        return None, "", "表内没有合计行"
    cols = prior_columns(ws, row, merged)
    if not cols:
        return None, rule, f"第 {row} 行（{row_label(ws, row, merged)}）所在表没有上年 / 期初列"
    keep = pick_column(cols, field)
    if len(keep) != 1:
        shown = "、".join(f"{get_column_letter(c)}「{'/'.join(cols[c])}」" for c in (keep or sorted(cols)))
        return None, rule, f"字段「{field}」对应的上年列不唯一或没有：{shown}"
    return f"{get_column_letter(keep[0])}{row}", rule, ""


# ── 生成 ─────────────────────────────────────────────────────────────────────────


def prev3_cells(mapping: dict) -> list[tuple[str, str, str, tuple[str, str, str]]]:
    """映射表全部块里的三参 PREV 格 ``(底稿, sheet, cell_ref, 参数)``；同一公式位出现两次 ⇒ 报错。"""
    out: list[tuple[str, str, str, tuple[str, str, str]]] = []
    seen: set[tuple[str, str, str]] = set()
    for block in mapping.get("mappings") or []:
        code, sheet = (block.get("wp_code") or "").strip(), (block.get("sheet") or "").strip()
        for cell in block.get("cells") or []:
            m = PREV3_RE.match(str(cell.get("formula") or ""))
            if m is None:
                continue
            key = (code, sheet, (cell.get("cell_ref") or "").strip())
            if key in seen:
                raise SystemExit(f"🔴 映射表同一公式位出现两次：{key}")
            seen.add(key)
            out.append((*key, m.groups()))
    return out


def _entry(code: str, sheet: str, ref: str, args: tuple[str, str, str], templates: dict[str, str]) -> tuple[dict, str]:
    """→（条目, 定不了的原因）；原因非空 = 需要人工裁定。"""
    base: dict[str, Any] = {"prev": list(args), "template": "", "cell": None, "header": [], "row_label": "",
                            "cell_is_formula": False, "action": "none", "rule": "", "reason": ""}
    manual = MANUAL.get((code, sheet, ref))
    found = _find_sheet(code, sheet)
    if found is None:
        names = [p.name for p, _ in _workbooks(code)]
        reason = (manual or {}).get("reason") or f"系统模板里没有该表（{'、'.join(names) or '无 xlsx 模板'}）"
        return {**base, "rule": "manual" if manual else "模板无该表", "reason": reason}, ""
    path, ws = found
    templates[_rel(path)] = _sha256(path)
    base["template"] = _rel(path)
    if manual is not None:
        if manual.get("action") == "none":
            return {**base, "rule": "manual", "reason": manual["reason"]}, ""
        cell, rule, why = manual["cell"], "manual", ""
    else:
        cell, rule, why = auto_locate(ws, args[2], "审定表" in sheet)
        if cell is None:
            return base, why
    header, label, is_formula = target_signature(ws, cell)
    if not has_prior_semantics(header, label):
        return base, f"{cell} 的列头「{'/'.join(header)}」与行标签「{label}」都没有上年 / 期初语义（或是比较列）"
    return {**base, "cell": cell, "header": list(header), "row_label": label, "cell_is_formula": is_formula,
            "action": "verify" if is_formula else "write", "rule": rule,
            "reason": (manual or {}).get("reason", "")}, ""


def has_prior_semantics(header: tuple[str, ...], label: str) -> bool:
    """目标格属上年 / 期初语义：列头是（非比较类的）上年列，或行标签写明上年（如「一、上年年末余额」）。"""
    if any(t in p for p in _body(header) for t in COMPARISON_TOKENS):
        return False
    return is_prior_header(header) or any(t in label for t in PRIOR_TOKENS)


def _conflicts(entries: dict[str, dict[str, dict[str, dict]]]) -> list[str]:
    """同一目标格被取数参数不同的两格占用 ⇒ 需人工裁定保留哪一格。"""
    by_target: dict[tuple[str, str, str], list[tuple[str, tuple]]] = {}
    for code, sheets in entries.items():
        for sheet, cells in sheets.items():
            for ref, e in cells.items():
                if e["action"] != "none":
                    by_target.setdefault((code, sheet, e["cell"]), []).append((ref, tuple(e["prev"])))
    return [f"{c} | {s} | {cell}：{refs}" for (c, s, cell), refs in sorted(by_target.items())
            if len({p for _, p in refs}) > 1]


def build() -> tuple[dict, list[str]]:
    mapping = json.loads(MAPPING.read_text(encoding="utf-8"))
    cells = prev3_cells(mapping)
    stale = sorted(set(MANUAL) - {(c, s, r) for c, s, r, _ in cells})
    if stale:
        raise SystemExit(f"🔴 MANUAL 里有映射表已没有的格（失效条目，删掉）：{stale}")
    missing_reason = [k for k, v in MANUAL.items() if not (v.get("reason") or "").strip()]
    if missing_reason:
        raise SystemExit(f"🔴 MANUAL 条目必须写原因：{missing_reason}")
    templates: dict[str, str] = {}
    entries: dict[str, dict[str, dict[str, dict]]] = {}
    unresolved: list[str] = []
    for code, sheet, ref, args in cells:
        entry, why = _entry(code, sheet, ref, args, templates)
        if why:
            unresolved.append(f"{code} | {sheet} | {ref} | PREV{args}：{why}")
            continue
        entries.setdefault(code, {}).setdefault(sheet, {})[ref] = entry
    unresolved += [f"同一目标格取数参数不同：{c}" for c in _conflicts(entries)]
    doc = {
        "_doc": "三参 PREV 预填逐格定位表（生成器 backend/scripts/gen/gen_prev3_prefill_locator.py，勿手改）",
        "source": "backend/data/prefill_formula_mapping.json",
        "templates": dict(sorted(templates.items())),
        "entries": {c: {s: dict(sorted(v.items())) for s, v in sorted(sheets.items())}
                    for c, sheets in sorted(entries.items())},
    }
    return doc, unresolved


def render(doc: dict) -> str:
    return json.dumps(doc, ensure_ascii=False, indent=1) + "\n"


def write_report(doc: dict, unresolved: list[str], path: Path) -> None:
    lines = ["# 三参 PREV 定位表审阅清单（不入库）", ""]
    for code, sheets in doc["entries"].items():
        for sheet, cells in sheets.items():
            for ref, e in cells.items():
                lines.append(f"{code} | {sheet} | {ref} | PREV{tuple(e['prev'])} → {e['action']} {e['cell'] or '-'} "
                             f"[{e['rule']}] 列头「{'/'.join(e['header'])}」行「{e['row_label']}」 {e['reason']}")
    lines += ["", f"## 待人工裁定 {len(unresolved)} 格", *unresolved]
    path.write_text("\n".join(lines) + "\n", encoding="utf-8")


def main(argv: list[str] | None = None) -> int:
    parser = argparse.ArgumentParser(description="生成三参 PREV 预填逐格定位表")
    parser.add_argument("--check", action="store_true", help="只比对，不写文件；漂移 exit 2")
    parser.add_argument("--report", type=Path, help="另写逐格审阅清单（不入库）")
    args = parser.parse_args(argv)
    doc, unresolved = build()
    if args.report:
        write_report(doc, unresolved, args.report)
    if unresolved:
        print(f"🔴 {len(unresolved)} 格自动规则定不了，需在 MANUAL 人工裁定：")
        print("\n".join(f"  {u}" for u in unresolved))
        return 1
    expected = render(doc)
    current = LOCATOR_FILE.read_text(encoding="utf-8") if LOCATOR_FILE.exists() else ""
    if args.check:
        if current != expected:
            print(f"[DRIFT] {LOCATOR_FILE.name} 与映射表 / 模板不同步，请重跑生成脚本")
            return 2
        print(f"[OK] {LOCATOR_FILE.name} 与映射表 / 模板同步")
        return 0
    if current == expected:
        print(f"(无变化) {LOCATOR_FILE.name}")
        return 0
    LOCATOR_FILE.write_text(expected, encoding="utf-8")
    print(f"[written] {LOCATOR_FILE.name}")
    return 0


if __name__ == "__main__":
    sys.exit(main())

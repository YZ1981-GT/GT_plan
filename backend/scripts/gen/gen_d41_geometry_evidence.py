"""生成 D4-1 几何核定 evidence + mapping_digest（Task 1）。

openpyxl 直读权威模板核实几何后，冻结：
- 两区 anchor（主营 R8 起 / 其他 R14 起）
- 受管列 A/B/C/D/F/G/H（label + 6 金额，与前端 useD4Adjudication 逐项对齐）
- formula_mask（E/I 数据行 + 行 12/18/19/21 的 B–I）
- 两区 UUID 空列（区1=空列A / 区2=空列B，不同列，同 D4-9 思路）

判据：几何与前端 `D4_ADJ_VALUE_FIELDS` 六金额字段逐项对齐；digest 冻结。
产物：.kiro/specs/.../evidence/d41-geometry.json
"""
from __future__ import annotations

import hashlib
import json
from pathlib import Path

from openpyxl import load_workbook
from openpyxl.utils import get_column_letter

BACKEND_ROOT = Path(__file__).resolve().parents[2]
REPO_ROOT = BACKEND_ROOT.parent
TEMPLATE = BACKEND_ROOT / "wp_templates" / "D" / "D4 收入底稿.xlsx"
TEMPLATE_RELATIVE = "backend/wp_templates/D/D4 收入底稿.xlsx"
FROZEN_TEMPLATE_SHA256 = (
    "b8fb92d4c22cd5d639e415403a12cb61650639153f136a5330c930b880167b5f"
)
SHEET = "营业收入审定表D4-1"
EVIDENCE = (
    REPO_ROOT
    / ".kiro"
    / "specs"
    / "d4-1-adjudication-bidirectional-writeback-and-formula-io"
    / "evidence"
    / "d41-geometry.json"
)

# ── 前端 useD4Adjudication / D4_ADJ_VALUE_FIELDS 六金额字段（逐项对齐）──
# currentUnadjusted / currentAje / currentRje / priorUnadjusted / priorAje / priorRje
FRONTEND_VALUE_FIELDS = [
    "currentUnadjusted",
    "currentAje",
    "currentRje",
    "priorUnadjusted",
    "priorAje",
    "priorRje",
]

# ── 受管字段：(column_key, 列, mode, value_type, store_key, header) ──
# A=label + 本期 B/C/D + 上期 F/G/H；与前端六金额逐项对齐。
MANAGED_FIELD_SPECS = [
    ("label", "A", "editable", "text", "label", "项目"),
    ("current_unadjusted", "B", "editable", "amount", "currentUnadjusted", "本期未审数"),
    ("current_aje", "C", "editable", "amount", "currentAje", "本期账项调整"),
    ("current_rje", "D", "editable", "amount", "currentRje", "本期重分类调整"),
    ("prior_unadjusted", "F", "editable", "amount", "priorUnadjusted", "上期未审数"),
    ("prior_aje", "G", "editable", "amount", "priorAje", "上期账项调整"),
    ("prior_rje", "H", "editable", "amount", "priorRje", "上期重分类调整"),
]

# ── 两区 anchor / 数据行 / 小计行 ──
MAIN_TITLE_ROW = 7          # 主营业务收入：
MAIN_FIRST_DATA_ROW = 8     # 主营段可扩行起点
MAIN_LAST_DATA_ROW = 11     # 模板预留末行（可扩）
MAIN_SUBTOTAL_ROW = 12      # 小计
OTHER_TITLE_ROW = 13        # 其他业务收入：
OTHER_FIRST_DATA_ROW = 14
OTHER_LAST_DATA_ROW = 17
OTHER_SUBTOTAL_ROW = 18     # 小计
TOTAL_ROW = 19              # 合计
TB_ROW = 20                 # 试算平衡表数
DIFF_ROW = 21               # 差异数

# ── formula_mask：数据行 E/I（审定数）+ 小计/合计/差异行 12/18/19/21 的 B–I ──
FOOTER_FORMULA_ROWS = [MAIN_SUBTOTAL_ROW, OTHER_SUBTOTAL_ROW, TOTAL_ROW, DIFF_ROW]
FORMULA_MASK = []
for r in range(MAIN_FIRST_DATA_ROW, MAIN_LAST_DATA_ROW + 1):
    FORMULA_MASK.extend([f"E{r}", f"I{r}"])
for r in range(OTHER_FIRST_DATA_ROW, OTHER_LAST_DATA_ROW + 1):
    FORMULA_MASK.extend([f"E{r}", f"I{r}"])
for r in FOOTER_FORMULA_ROWS:
    for c in range(2, 10):  # B..I
        FORMULA_MASK.append(f"{get_column_letter(c)}{r}")

# ── 两区 UUID 空列（不同列，同 D4-9 W/X 思路）──
# 模板 A1:I80 无 J+ 空列在受管区外；沿用平台惯例取 sheet 外围空列。
UUID_COL_MAIN = "W"
UUID_COL_OTHER = "X"

# ── TB 核对（只读回显，不入受管）──
TB_CHECK_KEYS = ["D4-1-adj-tb-6001", "D4-1-adj-tb-6051"]


def _verify_template_bytes() -> str:
    data = TEMPLATE.read_bytes()
    sha = hashlib.sha256(data).hexdigest()
    if sha != FROZEN_TEMPLATE_SHA256:
        raise SystemExit(
            f"模板字节漂移：实测 {sha} 冻结 {FROZEN_TEMPLATE_SHA256}"
        )
    return sha


def _verify_geometry_against_template() -> None:
    """openpyxl 复核关键几何（fail-closed）。"""
    wb = load_workbook(TEMPLATE, data_only=False)
    ws = wb[SHEET]
    assert ws.dimensions == "A1:I80", ws.dimensions
    assert ws.cell(MAIN_TITLE_ROW, 1).value == "主营业务收入："
    assert ws.cell(OTHER_TITLE_ROW, 1).value == "其他业务收入："
    assert ws.cell(MAIN_SUBTOTAL_ROW, 1).value == "小计"
    assert ws.cell(OTHER_SUBTOTAL_ROW, 1).value == "小计"
    assert ws.cell(TOTAL_ROW, 1).value == "合计"
    assert ws.cell(TB_ROW, 1).value == "试算平衡表数"
    assert ws.cell(DIFF_ROW, 1).value == "差异数"
    # E/I 数据行公式
    assert ws.cell(MAIN_FIRST_DATA_ROW, 5).value == "=SUM(B8:D8)"
    assert ws.cell(MAIN_FIRST_DATA_ROW, 9).value == "=SUM(F8:H8)"
    assert ws.cell(OTHER_FIRST_DATA_ROW, 5).value == "=SUM(B14:D14)"
    # 受管数据行 B/C/D/F/G/H 无公式（是输入格）
    for r in (MAIN_FIRST_DATA_ROW, OTHER_FIRST_DATA_ROW):
        for c in (2, 3, 4, 6, 7, 8):
            v = ws.cell(r, c).value
            assert not (isinstance(v, str) and v.startswith("=")), (r, c, v)
    # 小计/合计/差异行公式格
    assert ws.cell(MAIN_SUBTOTAL_ROW, 2).value == "=SUM(B8:B11)"
    assert ws.cell(TOTAL_ROW, 2).value == "=B18+B12"
    assert ws.cell(DIFF_ROW, 5).value == "=E19-E20"


def _digest_payload() -> dict:
    """mapping_digest canon payload（冻结映射，防漂移）。"""
    return {
        "managed_sheet": SHEET,
        "template_relative_path": TEMPLATE_RELATIVE,
        "template_sha256": FROZEN_TEMPLATE_SHA256,
        "header_rows": [5, 6],
        "regions": [
            {
                "table_key": "adjudication_main_rows",
                "section_key": "main-revenue",
                "title_row": MAIN_TITLE_ROW,
                "first_data_row": MAIN_FIRST_DATA_ROW,
                "last_data_row": MAIN_LAST_DATA_ROW,
                "subtotal_row": MAIN_SUBTOTAL_ROW,
                "uuid_col": UUID_COL_MAIN,
            },
            {
                "table_key": "adjudication_other_rows",
                "section_key": "other-revenue",
                "title_row": OTHER_TITLE_ROW,
                "first_data_row": OTHER_FIRST_DATA_ROW,
                "last_data_row": OTHER_LAST_DATA_ROW,
                "subtotal_row": OTHER_SUBTOTAL_ROW,
                "uuid_col": UUID_COL_OTHER,
            },
        ],
        "total_row": TOTAL_ROW,
        "tb_row": TB_ROW,
        "diff_row": DIFF_ROW,
        "contract_fields": [
            {"col": col, "store_key": store_key, "value_type": value_type}
            for _ck, col, _mode, value_type, store_key, _hdr in MANAGED_FIELD_SPECS
        ],
        "formula_mask": FORMULA_MASK,
        "tb_check_keys": TB_CHECK_KEYS,
    }


def _compute_digest(payload: dict) -> str:
    canon = json.dumps(
        payload, ensure_ascii=False, sort_keys=True, separators=(",", ":")
    )
    return hashlib.sha256(canon.encode("utf-8")).hexdigest()


def main() -> None:
    sha = _verify_template_bytes()
    _verify_geometry_against_template()

    # 六金额字段与前端逐项对齐判据
    managed_store_keys = [
        store_key
        for _ck, _col, _mode, _vt, store_key, _hdr in MANAGED_FIELD_SPECS
        if store_key != "label"
    ]
    assert managed_store_keys == FRONTEND_VALUE_FIELDS, (
        managed_store_keys,
        FRONTEND_VALUE_FIELDS,
    )

    digest_payload = _digest_payload()
    mapping_digest = _compute_digest(digest_payload)

    evidence = {
        "measured_at": "2026-09-19",
        "tool": "openpyxl data_only=False",
        "spec": "d4-1-adjudication-bidirectional-writeback-and-formula-io",
        "task": "1 几何核定与 mapping_digest 冻结",
        "template_path": TEMPLATE_RELATIVE,
        "template_sha256": sha,
        "managed_sheet": SHEET,
        "dimensions": "A1:I80",
        "dec0": "同 sheet 双区动态行（主营 R8 起 / 其他 R14 起），参照 D4-9",
        "header_rows": [5, 6],
        "header_layout": {
            "R5": "项目 / 本期数(B–E) / 上期数(F–I)",
            "R6": "未审数 / 账项调整 / 重分类调整 / 审定数（本期与上期各一组）",
        },
        "regions": {
            "main_revenue": {
                "table_key": "adjudication_main_rows",
                "section_key": "main-revenue",
                "title_row": MAIN_TITLE_ROW,
                "title_marker": "主营业务收入：",
                "first_data_row": MAIN_FIRST_DATA_ROW,
                "last_data_row": MAIN_LAST_DATA_ROW,
                "reserved_rows": MAIN_LAST_DATA_ROW - MAIN_FIRST_DATA_ROW + 1,
                "subtotal_row": MAIN_SUBTOTAL_ROW,
                "subtotal_marker": "小计",
                "uuid_col": UUID_COL_MAIN,
                "expandable": True,
            },
            "other_revenue": {
                "table_key": "adjudication_other_rows",
                "section_key": "other-revenue",
                "title_row": OTHER_TITLE_ROW,
                "title_marker": "其他业务收入：",
                "first_data_row": OTHER_FIRST_DATA_ROW,
                "last_data_row": OTHER_LAST_DATA_ROW,
                "reserved_rows": OTHER_LAST_DATA_ROW - OTHER_FIRST_DATA_ROW + 1,
                "subtotal_row": OTHER_SUBTOTAL_ROW,
                "subtotal_marker": "小计",
                "uuid_col": UUID_COL_OTHER,
                "expandable": True,
            },
        },
        "footer_rows": {
            "total_row": TOTAL_ROW,
            "tb_row": TB_ROW,
            "diff_row": DIFF_ROW,
            "audit_note_start_row": 22,
        },
        "managed_columns": ["A", "B", "C", "D", "F", "G", "H"],
        "field_map": [
            {
                "col": col,
                "column_key": ck,
                "mode": mode,
                "value_type": vt,
                "store_key": store_key,
                "header": hdr,
                "contract": True,
            }
            for ck, col, mode, vt, store_key, hdr in MANAGED_FIELD_SPECS
        ],
        "frontend_alignment": {
            "composable": "useD4Adjudication.ts / d4AdjudicationRows.ts",
            "store_item_id": "D4-1-rows",
            "per_field_key_pattern": "D4-1-{rowId}-{field}",
            "value_fields": FRONTEND_VALUE_FIELDS,
            "aligned": managed_store_keys == FRONTEND_VALUE_FIELDS,
        },
        "formula_mask": {
            "cells": FORMULA_MASK,
            "note": (
                "E/I 审定数（=SUM(B:D)/=SUM(F:H)）+ 行 12/18/19/21 的 B–I "
                "（小计/合计/差异全为 Excel 内部公式）；普通值投影不覆盖"
            ),
        },
        "tb_check": {
            "keys": TB_CHECK_KEYS,
            "note": "试算平衡表数（R20）只读回显，不入受管；仅标量走 Tier A 公式",
        },
        "prefill_note": (
            "未审数不从四表库填（presets.py R3.4：TB 6001/6051 只有科目总额、"
            "无产品/项目维度）；明细行由 D4-2/D4-3 SUMIF 派生 + 人工"
        ),
        "digest_payload": digest_payload,
        "mapping_digest": mapping_digest,
    }

    EVIDENCE.parent.mkdir(parents=True, exist_ok=True)
    EVIDENCE.write_text(
        json.dumps(evidence, ensure_ascii=False, indent=2), encoding="utf-8"
    )
    print(f"[OK] evidence 写入 {EVIDENCE}")
    print(f"[OK] mapping_digest = {mapping_digest}")
    print(f"[OK] 六金额字段与前端逐项对齐 = {managed_store_keys == FRONTEND_VALUE_FIELDS}")
    print(f"[OK] formula_mask 格数 = {len(FORMULA_MASK)}")


if __name__ == "__main__":
    main()

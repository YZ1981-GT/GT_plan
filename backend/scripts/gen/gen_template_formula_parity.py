# -*- coding: utf-8 -*-
r"""生成「前端 ↔ 模板公式」对拍清单（spec oo-writeback-template-parity-closure R2.4，design §3.3）。

读每份生产契约（``backend/data/workpaper_sync_contracts/*.json``，不含 ``.candidate.``）里的**行表**
（有 ``row_identity`` 的表），对每个 ``mode=formula`` 的字段取权威模板数据首行那一格的公式，
把同行单元格引用换成同一张表的 store 字段名，写入
``audit-platform/frontend/src/components/workpaper/sync/__tests__/templateFormulaParity.generated.json``。
前端 ``templateFormulaParity.spec.ts`` 用它与前端行计算逐值对拍。

换不成字段名表达式的公式标 ``skip`` 并写原因，不猜：引用他行 / 他表 / 定义名、引用了非受管且非空的列、
函数不在支持范围（只支持 SUM / IF / AND / OR / NOT / ROUND / ABS / MAX / MIN / IFERROR）、
字符串拼接、数组公式。

输出不含时间戳，内容不变则字节不变（``--check`` 才成立）。

用法（backend 目录）：
  生成:  ..\.venv\Scripts\python.exe scripts/gen/gen_template_formula_parity.py
  校验:  ..\.venv\Scripts\python.exe scripts/gen/gen_template_formula_parity.py --check   （漂移 exit 2）
  其他树: --backend-root <某 worktree 的 backend 目录>  --out <输出路径>
"""
from __future__ import annotations

import io
import json
import re
import sys
import warnings
from dataclasses import dataclass, field
from pathlib import Path

ROOT = Path(__file__).resolve().parents[2]  # backend/
OUT_REL = Path("audit-platform/frontend/src/components/workpaper/sync/__tests__/templateFormulaParity.generated.json")

SUPPORTED_FUNCS = frozenset({"SUM", "IF", "AND", "OR", "NOT", "ROUND", "ABS", "MAX", "MIN", "IFERROR"})
#: 取值为数的字段类型（对拍时随机生成输入）
NUMERIC_TYPES = frozenset({"amount", "number", "ratio", "percent", "integer", "decimal", "rate"})
_REF = re.compile(r"^(\$?)([A-Z]{1,3})(\$?)(\d+)$")
_SRC = re.compile(r"^源xlsx!(?P<sheet>.+)!(?P<cell>\$?[A-Z]{1,3}\$?\d+)$")


class Skip(Exception):
    """公式换不成同行字段表达式。"""


@dataclass
class Column:
    field: str
    mode: str
    value_type: str
    header: str


@dataclass
class TableCtx:
    contract: str
    sheet: str
    store_item: str
    table_key: str
    data_row: int
    columns: dict[str, Column] = field(default_factory=dict)
    #: 模板里数据首行无值无公式、表头区也没有文字的列（间隔列）⇒ 引用按 0
    blank_columns: set[str] = field(default_factory=set)
    literals: set[str] = field(default_factory=set)
    ws: object = None
    inlining: tuple[str, ...] = ()  # 正在展开的非受管公式列（防循环）


def _field_name(pointer: str) -> str:
    """行内字段路径：``/rows/{row_uuid}/aging/within1`` → ``aging/within1``（只取末段会撞名）。"""
    tail = pointer.split("{row_uuid}/", 1)[-1] if "{row_uuid}/" in pointer else pointer.rsplit("/", 1)[-1]
    return tail.strip("/")


def _split_ref(text: str) -> tuple[str, int, bool]:
    m = _REF.match(text)
    if m is None:
        raise Skip(f"引用定义名或整列 {text}")
    return m.group(2), int(m.group(4)), bool(m.group(3))


def _operand(ctx: TableCtx, col: str) -> str:
    column = ctx.columns.get(col)
    if column is not None:
        return column.field
    if col in ctx.blank_columns:
        return "0"
    value = ctx.ws[f"{col}{ctx.data_row}"].value if ctx.ws is not None else None
    # 非受管的公式列（如「本期增加」小计列）：同行展开它的公式；只有常数的列按常数
    if isinstance(value, str) and value.startswith("=") and col not in ctx.inlining and len(ctx.inlining) < 4:
        saved = ctx.inlining
        ctx.inlining = saved + (col,)
        try:
            return f"( {convert(value, ctx)} )"
        finally:
            ctx.inlining = saved
    if isinstance(value, (int, float)) and not isinstance(value, bool):
        return repr(value)
    raise Skip(f"引用非受管列 {col}")


def _cell_operand(ctx: TableCtx, text: str, in_range_func: bool) -> str:
    if "!" in text:
        raise Skip(f"引用他表 {text}")
    if ":" in text:
        if not in_range_func:
            raise Skip(f"区间不在 SUM / MAX / MIN 内 {text}")
        a, b = text.split(":", 1)
        ca, ra, la = _split_ref(a)
        cb, rb, lb = _split_ref(b)
        if ra != ctx.data_row or rb != ctx.data_row or la or lb:
            raise Skip(f"引用他行 {text}")
        from openpyxl.utils import column_index_from_string, get_column_letter

        lo, hi = sorted((column_index_from_string(ca), column_index_from_string(cb)))
        return ", ".join(_operand(ctx, get_column_letter(i)) for i in range(lo, hi + 1))
    col, row, locked = _split_ref(text)
    if row != ctx.data_row or locked:
        raise Skip(f"引用他行 {text}")
    return _operand(ctx, col)


def convert(formula: str, ctx: TableCtx) -> str:
    """模板公式 → 字段名表达式（token 以空格分隔）；换不成抛 ``Skip``。"""
    from openpyxl.formula.tokenizer import Token, Tokenizer

    out: list[str] = []
    stack: list[str] = []  # 函数名 / "(" 栈，判断区间是否在 SUM / MAX / MIN 内
    for tok in Tokenizer(formula).items:
        if tok.type == Token.WSPACE:
            continue
        if tok.type == Token.OPERAND:
            if tok.subtype == Token.RANGE:
                in_range_func = bool(stack) and stack[-1] in {"SUM", "MAX", "MIN"}
                out.append(_cell_operand(ctx, tok.value, in_range_func))
            elif tok.subtype == Token.NUMBER:
                out.append(tok.value)
            elif tok.subtype == Token.TEXT:
                ctx.literals.add(tok.value[1:-1].replace('""', '"'))
                out.append(tok.value)
            elif tok.subtype == Token.LOGICAL:
                out.append(tok.value.upper())
            else:
                raise Skip(f"不支持的操作数 {tok.value}")
        elif tok.type == Token.FUNC:
            if tok.subtype == Token.OPEN:
                name = tok.value[:-1].upper().removeprefix("_XLFN.")
                if name not in SUPPORTED_FUNCS:
                    raise Skip(f"函数 {name} 不在支持范围")
                stack.append(name)
                out.append(f"{name}(")
            else:
                stack.pop()
                out.append(")")
        elif tok.type == Token.PAREN:
            if tok.subtype == Token.OPEN:
                stack.append("(")
            else:
                stack.pop()
            out.append(tok.value)
        elif tok.type == Token.SEP:
            if tok.subtype != Token.ARG:
                raise Skip("数组常量")
            out.append(",")
        elif tok.type == Token.OP_IN:
            if tok.value == "&":
                raise Skip("字符串拼接")
            out.append(tok.value)
        elif tok.type == Token.OP_PRE:
            out.append(tok.value)
        elif tok.type == Token.OP_POST:
            raise Skip(f"后缀运算 {tok.value}")
        else:
            raise Skip(f"不支持的记号 {tok.type}")
    return " ".join(out)


def _load_sheet(book, sheet: str):
    if sheet not in book.sheetnames:
        raise Skip(f"模板无 sheet {sheet}")
    return book[sheet]


def collect(backend_root: Path) -> dict:
    import openpyxl

    warnings.filterwarnings("ignore")
    contracts_dir = backend_root / "data" / "workpaper_sync_contracts"
    entries: list[dict] = []
    items: dict[str, dict] = {}
    books: dict[str, object] = {}
    for path in sorted(contracts_dir.glob("*.json")):
        if ".candidate." in path.name:
            continue
        contract = json.loads(path.read_text(encoding="utf-8"))
        cid = contract.get("contract_id") or path.stem
        rel = (contract.get("template") or {}).get("relative_path")
        if not rel:
            continue
        if rel not in books:
            data = (backend_root / "wp_templates" / rel).read_bytes()
            books[rel] = openpyxl.load_workbook(io.BytesIO(data), data_only=False)
        book = books[rel]
        for sheet in contract.get("sheets", []):
            for table in sheet.get("tables", []) or []:
                if "row_identity" not in table:
                    continue
                _collect_table(cid, sheet["excel_name"], table, book, entries, items)
    for book in books.values():
        book.close()
    entries.sort(key=lambda e: (e["store_item"], e["contract"], e["field"]))
    ok = sum(1 for e in entries if e["expr"] is not None)
    return {
        "_source": "backend/scripts/gen/gen_template_formula_parity.py（勿手改；重生成后 --check 无漂移）",
        "summary": {"formula_fields": len(entries), "ok": ok, "skip": len(entries) - ok},
        "items": dict(sorted(items.items())),
        "entries": entries,
    }


def _collect_table(cid: str, sheet_name: str, table: dict, book, entries: list, items: dict) -> None:
    fields = table.get("fields", [])
    formula_fields = [f for f in fields if f.get("mode") == "formula"]
    if not formula_fields:
        return
    first = _SRC.match(formula_fields[0]["source_ref"])
    data_row = int(re.sub(r"[^0-9]", "", first.group("cell"))) if first else 0
    store_item = formula_fields[0].get("store_item_id") or ""
    ctx = TableCtx(cid, sheet_name, store_item, table.get("table_key", ""), data_row)
    for f in fields:
        col = (f.get("cell") or {}).get("column")
        if col:
            ctx.columns[col] = Column(_field_name(f["json_pointer"]), f.get("mode", ""),
                                      f.get("value_type", ""), f.get("header_text", ""))
    try:
        ws = _load_sheet(book, sheet_name)
    except Skip as exc:
        ws, sheet_skip = None, str(exc)
    else:
        from openpyxl.utils import get_column_letter

        ctx.ws = ws
        sheet_skip = None
        header_top = max(1, data_row - int(table.get("header_rows") or 1) - 2)
        headed: set[int] = set()  # 表头区有文字的列（合并表头按整块算）
        for rng in ws.merged_cells.ranges:
            if rng.max_row >= header_top and rng.min_row < data_row and \
                    ws.cell(rng.min_row, rng.min_col).value not in (None, ""):
                headed.update(range(rng.min_col, rng.max_col + 1))
        for ci in range(1, ws.max_column + 1):
            letter = get_column_letter(ci)
            if letter in ctx.columns or ws.cell(data_row, ci).value is not None:
                continue
            # 表头区有文字 ⇒ 是没受管的录入列（OO 里可填），不能按 0
            if ci not in headed and all(ws.cell(r, ci).value in (None, "") for r in range(header_top, data_row)):
                ctx.blank_columns.add(letter)
    key = store_item or f"<未接通>{cid}/{ctx.table_key}"
    item = items.setdefault(key, {"contract": cid, "tables": [], "inputs": [], "input_types": {},
                                  "date_inputs": [], "text_inputs": [], "auto_fields": [], "literals": []})
    item["tables"].append(f"{sheet_name}!{ctx.table_key}@{data_row}")
    # 只有 editable 是录入列；auto_source / cross_sheet 由系统带出（受保护），既不随机录入也不按录入核对
    editable = [c for c in ctx.columns.values() if c.mode == "editable"]
    item["auto_fields"] = sorted(set(item["auto_fields"]) | {
        c.field for c in ctx.columns.values() if c.mode not in ("editable", "formula")})
    for slot, cols in (("inputs", [c for c in editable if c.value_type in NUMERIC_TYPES]),
                       ("date_inputs", [c for c in editable if c.value_type == "date"]),
                       ("text_inputs", [c for c in editable if c.value_type not in NUMERIC_TYPES | {"date"}])):
        item[slot] = sorted(set(item[slot]) | {c.field for c in cols})
    # 数值录入列的取值类型（对拍按类型取随机值：比例 0~1、整数、金额）
    types = dict(item["input_types"])
    for c in editable:
        if c.value_type in NUMERIC_TYPES:
            types.setdefault(c.field, c.value_type)
    item["input_types"] = dict(sorted(types.items()))
    for f in formula_fields:
        m = _SRC.match(f["source_ref"])
        cell = m.group("cell") if m else f["source_ref"]
        entry = {"store_item": key, "contract": cid, "sheet": sheet_name, "cell": cell,
                 "field": _field_name(f["json_pointer"]), "header": f.get("header_text", ""),
                 "formula": None, "expr": None, "skip": None}
        try:
            if not store_item or store_item == "placeholder":
                raise Skip(f"store item 未接通（{store_item or '空'}，R3）")
            if sheet_skip:
                raise Skip(sheet_skip)
            value = ws[cell].value
            if value is None or not isinstance(value, str) or not value.startswith("="):
                entry["formula"] = None if value is None else str(value)
                raise Skip("模板数据首行该格不是公式")
            entry["formula"] = value
            entry["expr"] = convert(value, ctx)
        except Skip as exc:
            entry["skip"] = str(exc)
        except Exception as exc:  # 数组公式等 openpyxl 对象
            entry["skip"] = f"无法解析：{type(exc).__name__}"
        entries.append(entry)
    item["literals"] = sorted(set(item["literals"]) | ctx.literals)


def render(payload: dict) -> str:
    return json.dumps(payload, ensure_ascii=False, indent=1, sort_keys=False) + "\n"


def main(argv: list[str]) -> int:
    backend_root = Path(argv[argv.index("--backend-root") + 1]) if "--backend-root" in argv else ROOT
    out = Path(argv[argv.index("--out") + 1]) if "--out" in argv else backend_root.parent / OUT_REL
    text = render(collect(backend_root))
    if "--check" in argv:
        current = out.read_text(encoding="utf-8") if out.exists() else ""
        if current != text:
            print(f"[DRIFT] {out} 与契约 / 模板现算不一致，请重跑生成器")
            return 2
        print(f"[ok] {out} 无漂移")
        return 0
    out.parent.mkdir(parents=True, exist_ok=True)
    if not out.exists() or out.read_text(encoding="utf-8") != text:
        out.write_text(text, encoding="utf-8", newline="\n")
        print(f"[write] {out}")
    else:
        print(f"[same] {out}")
    return 0


if __name__ == "__main__":
    sys.stdout.reconfigure(encoding="utf-8", errors="replace")
    raise SystemExit(main(sys.argv[1:]))

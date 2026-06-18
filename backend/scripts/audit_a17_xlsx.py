"""Deep-read A17 程序表 + A17-5-1~5 核对表 xlsx; produce a17_xlsx_audit.json.

PRE-4-0 / X-A17：为 checklist_xlsx_parser（PRE-4-1/4-2）提供列映射权威来源。

产出范围（6 个 xlsx）：
  - A17 重大事项概要程序表（a-program-console，5 步）
  - A17-5-1 财报审计 / A17-5-2 内控审计 / A17-5-3 IPO / A17-5-4 新三板 / A17-5-5 函证
    （checklist-table，两节结构：一、审计目标[无结论列] + 二、核对程序[结论/索引列]）

item_id 注册表对齐 persistence.md：`A17-5-{x}-{seq}`（核对程序节内顺序号，零填充）。

用法：
    python scripts/audit_a17_xlsx.py            # 审计全部 6 文件
    python scripts/audit_a17_xlsx.py A17-5-1    # 仅指定 wp_code
"""
from __future__ import annotations

import argparse
import json
import re
from pathlib import Path

import openpyxl
from openpyxl.utils import get_column_letter

ROOT = Path(__file__).resolve().parents[2]
TEMPLATES = ROOT / "backend" / "wp_templates" / "A"
OUT = ROOT / "backend" / "data" / "a17_xlsx_audit.json"
PROC_JSON = ROOT / "backend" / "data" / "procedure_table_templates.json"

# wp_code -> 物理文件名前缀（A17-5-x 文件名空格不规则，用 startswith 容错）
A17_WP_CODES = ["A17", "A17-5-1", "A17-5-2", "A17-5-3", "A17-5-4", "A17-5-5"]


def cell_str(v) -> str | None:
    if v is None:
        return None
    s = str(v).strip()
    return s if s else None


def resolve_template_path(wp_code: str) -> Path:
    """按 wp_code 前缀匹配物理 xlsx；A17 程序表须排除 A17-x 子码。"""
    if wp_code == "A17":
        matches = [
            p
            for p in TEMPLATES.iterdir()
            if p.suffix.lower() == ".xlsx"
            and re.match(r"^A17\s", p.name)
            and not re.match(r"^A17-", p.name)
        ]
    else:
        # A17-5-1 等；文件名形如 "A17-5-1 审计工作完成核对表（…）.xlsx"，含双空格变体
        norm = wp_code
        matches = [
            p
            for p in TEMPLATES.iterdir()
            if p.suffix.lower() == ".xlsx"
            and p.name.replace("  ", " ").startswith(f"{norm} ")
        ]
    if not matches:
        raise FileNotFoundError(f"No xlsx template for {wp_code} under {TEMPLATES}")
    return sorted(matches)[0]


# ---------------------------------------------------------------------------
# A17 程序表（a-program-console）
# ---------------------------------------------------------------------------

def audit_program_sheet(ws, sn: str) -> dict:
    merged = [str(m) for m in ws.merged_cells.ranges]
    # 表头行：含「序号」「程序」「索引号」
    hr = 5
    for r in range(1, min(12, ws.max_row + 1)):
        texts = " ".join(cell_str(ws.cell(r, c).value) or "" for c in range(1, ws.max_column + 1))
        if "序号" in texts and "程序" in texts:
            hr = r
            break
    cols = []
    role_cols: dict[str, str] = {}
    for c in range(1, ws.max_column + 1):
        v = cell_str(ws.cell(hr, c).value)
        if not v:
            continue
        letter = get_column_letter(c)
        cols.append({"col": letter, "label": v})
        if v == "序号":
            role_cols["seq"] = letter
        elif v == "程序":
            role_cols["procedure"] = letter
        elif "是否适用" in v:
            role_cols["applicable"] = letter
        elif "执行人" in v:
            role_cols["performer"] = letter
        elif "执行情况" in v:
            role_cols["execution_note"] = letter
        elif v == "索引号":
            role_cols["ref_index"] = letter

    ref_ci = (
        openpyxl.utils.column_index_from_string(role_cols["ref_index"])
        if "ref_index" in role_cols
        else None
    )
    rows = []
    for r in range(hr + 1, ws.max_row + 1):
        a = ws.cell(r, 1).value
        b = cell_str(ws.cell(r, 2).value)
        if a is None and not b:
            continue
        seq_str = None
        if a is not None:
            seq_str = str(int(a)) if isinstance(a, (int, float)) and float(a).is_integer() else str(a).strip()
        if seq_str and seq_str.startswith(("提示", "说明")):
            continue
        entry = {
            "row": r,
            "seq": seq_str,
            "procedure_preview": (b[:150] if b else None),
        }
        if ref_ci:
            ref = cell_str(ws.cell(r, ref_ci).value)
            if ref:
                entry["ref_index"] = ref.replace("\n", ",").replace("、", ",")
        rows.append(entry)
    return {
        "name": sn,
        "component_type": "a-program-console",
        "header_row": hr,
        "max_row": ws.max_row,
        "max_col": ws.max_column,
        "columns": cols,
        "column_roles": role_cols,
        "data_rows": len(rows),
        "merged_ranges_count": len(merged),
        "procedure_rows": rows,
    }


def audit_program_table(path: Path, wp_code: str) -> dict:
    wb = openpyxl.load_workbook(path, data_only=True, read_only=False)
    sheets = []
    for sn in wb.sheetnames:
        if sn == "GT_Custom":
            sheets.append({"name": sn, "component_type": "skip"})
            continue
        sheets.append(audit_program_sheet(wb[sn], sn))
    wb.close()
    return {
        "wp_code": wp_code,
        "filename": path.name,
        "runtime": "a-program-console",
        "sheets": sheets,
        "parser_ready": True,
        "notes": "A17 重大事项概要程序表，5 步：seq1→A17-1 / seq2→A17-3 / seq3→A17-4 / "
                 "seq4→A17-2 / seq5→A17-6。A17-5 核对表不在程序表内（独立底稿）。",
    }


# ---------------------------------------------------------------------------
# A17-5-x 核对表（checklist-table，两节结构）
# ---------------------------------------------------------------------------

_OBJECTIVE_HDR_RE = re.compile(r"^一[、.\s]*审计目标")
_PROCEDURE_HDR_RE = re.compile(r"^二[、.\s]*核对程序")


def _detect_procedure_column_roles(ws, hr: int) -> tuple[list[dict], dict[str, str]]:
    """从「二、核对程序」表头行提取列标签与语义角色（结论/索引列位置因文件而异）。"""
    cols: list[dict] = []
    roles: dict[str, str] = {}
    for c in range(1, ws.max_column + 1):
        v = cell_str(ws.cell(hr, c).value)
        if not v:
            continue
        letter = get_column_letter(c)
        cols.append({"col": letter, "label": v})
        if "核对结果" in v:
            roles["conclusion"] = letter
        elif "工作底稿索引号" in v:
            roles["work_paper_index"] = letter  # 用户填写
        elif "索引号参考" in v or "参考底稿索引号" in v:
            roles["index_reference"] = letter  # 预置参考
        elif "披露章节" in v:
            roles["disclosure_chapter"] = letter
        elif v == "备注":
            # 可能出现两次（结论旁 + 索引旁）；记第一个为 remark
            roles.setdefault("remark", letter)
    return cols, roles


def audit_checklist_sheet(ws, sn: str, wp_code: str) -> dict:
    merged = [str(m) for m in ws.merged_cells.ranges]

    # 1) 定位两节锚点行
    objective_row = None
    procedure_row = None
    for r in range(1, ws.max_row + 1):
        a = cell_str(ws.cell(r, 1).value)
        if not a:
            continue
        if objective_row is None and _OBJECTIVE_HDR_RE.match(a):
            objective_row = r
        elif procedure_row is None and _PROCEDURE_HDR_RE.match(a):
            procedure_row = r
            break

    # 2) 一、审计目标 条目（无结论列，仅 seq + 文本）
    objectives = []
    if objective_row is not None:
        end = procedure_row if procedure_row else ws.max_row + 1
        for r in range(objective_row + 1, end):
            a = ws.cell(r, 1).value
            b = cell_str(ws.cell(r, 2).value)
            if a is None and not b:
                continue
            seq = str(int(a)) if isinstance(a, (int, float)) and float(a).is_integer() else cell_str(a)
            objectives.append({
                "row": r,
                "seq": seq,
                "text_preview": (b[:150] if b else None),
            })

    # 3) 二、核对程序 列角色 + 条目（带结论/索引列）
    procedure_cols: list[dict] = []
    procedure_roles: dict[str, str] = {}
    procedures = []
    if procedure_row is not None:
        procedure_cols, procedure_roles = _detect_procedure_column_roles(ws, procedure_row)
        concl_ci = (
            openpyxl.utils.column_index_from_string(procedure_roles["conclusion"])
            if "conclusion" in procedure_roles
            else None
        )
        wp_idx_ci = (
            openpyxl.utils.column_index_from_string(procedure_roles["work_paper_index"])
            if "work_paper_index" in procedure_roles
            else None
        )
        ref_ci = (
            openpyxl.utils.column_index_from_string(procedure_roles["index_reference"])
            if "index_reference" in procedure_roles
            else None
        )
        seq_no = 0
        for r in range(procedure_row + 1, ws.max_row + 1):
            a = ws.cell(r, 1).value
            b = cell_str(ws.cell(r, 2).value)
            a_str = cell_str(a)
            if a_str and a_str.startswith(("提示", "说明", "注：")):
                continue
            if a is None and not b:
                continue
            if a_str is None and not b:
                continue
            seq_no += 1
            seq_label = str(int(a)) if isinstance(a, (int, float)) and float(a).is_integer() else a_str
            entry = {
                "row": r,
                "seq_no": seq_no,
                "item_id": f"{wp_code}-{seq_no:03d}",
                "seq_label": seq_label,
                "content_preview": (b[:150] if b else None),
            }
            if ref_ci:
                preset = cell_str(ws.cell(r, ref_ci).value)
                if preset:
                    entry["preset_index_reference"] = preset
            # 预置参考索引可能直接写在 work_paper_index 列（A17-5-1 中 G=底稿索引号参考实为预置）
            if wp_idx_ci and "index_reference" not in procedure_roles:
                preset = cell_str(ws.cell(r, wp_idx_ci).value)
                if preset:
                    entry["preset_index_reference"] = preset
            procedures.append(entry)

    return {
        "name": sn,
        "component_type": "checklist-table",
        "max_row": ws.max_row,
        "max_col": ws.max_column,
        "objective_section": {
            "anchor_row": objective_row,
            "title": "一、审计目标",
            "has_conclusion": False,
            "item_count": len(objectives),
            "items": objectives,
        },
        "procedure_section": {
            "anchor_row": procedure_row,
            "title": "二、核对程序",
            "has_conclusion": True,
            "columns": procedure_cols,
            "column_roles": procedure_roles,
            "item_count": len(procedures),
            "item_id_pattern": f"{wp_code}-{{seq:03d}}",
            "items": procedures,
        },
        "merged_ranges_count": len(merged),
        "notes": "两节结构：审计目标（无结论列，纯说明）+ 核对程序（结论[是/否/不适用]+"
                 "工作底稿索引号(填)+预置参考索引+备注）。parser 须区分两节类型。",
    }


def audit_checklist(path: Path, wp_code: str) -> dict:
    wb = openpyxl.load_workbook(path, data_only=True, read_only=False)
    sheets = []
    for sn in wb.sheetnames:
        if sn == "GT_Custom":
            sheets.append({"name": sn, "component_type": "skip"})
            continue
        sheets.append(audit_checklist_sheet(wb[sn], sn, wp_code))
    wb.close()
    main_sheet = next((s for s in sheets if s.get("component_type") == "checklist-table"), None)
    proc_count = main_sheet["procedure_section"]["item_count"] if main_sheet else 0
    return {
        "wp_code": wp_code,
        "filename": path.name,
        "runtime": "checklist-table",
        "sheets": sheets,
        "parser_ready": True,
        "procedure_item_count": proc_count,
        "notes": f"A17-5-x 审计工作完成核对表；核对程序 {proc_count} 条；"
                 f"item_id 模式 {wp_code}-{{seq:03d}}（对齐 persistence.md A17-5-{{x}}-{{seq}}）。",
    }


# ---------------------------------------------------------------------------
# dispatch + main
# ---------------------------------------------------------------------------

def audit_wp(wp_code: str) -> dict:
    path = resolve_template_path(wp_code)
    if wp_code == "A17":
        return audit_program_table(path, wp_code)
    if wp_code.startswith("A17-5-"):
        return audit_checklist(path, wp_code)
    raise NotImplementedError(f"audit handler not implemented for {wp_code}")


def merge_audit(entry: dict) -> None:
    existing: list[dict] = []
    if OUT.exists():
        existing = json.loads(OUT.read_text(encoding="utf-8"))
    existing = [e for e in existing if e.get("wp_code") != entry["wp_code"]]
    existing.append(entry)
    existing.sort(key=lambda x: x["wp_code"])
    OUT.write_text(json.dumps(existing, ensure_ascii=False, indent=2), encoding="utf-8")


def main() -> None:
    parser = argparse.ArgumentParser(description="Audit A17 程序表 + A17-5-1~5 核对表 xlsx")
    parser.add_argument(
        "wp_codes",
        nargs="*",
        default=A17_WP_CODES,
        help=f"默认全部：{' '.join(A17_WP_CODES)}",
    )
    args = parser.parse_args()
    missing: list[str] = []
    for code in args.wp_codes:
        try:
            entry = audit_wp(code)
        except FileNotFoundError as exc:
            missing.append(code)
            print(f"[MISSING] {code}: {exc}")
            merge_audit({
                "wp_code": code,
                "filename": None,
                "runtime": None,
                "sheets": [],
                "parser_ready": False,
                "data_blocked": True,
                "notes": f"物理文件缺失：{exc}",
            })
            continue
        merge_audit(entry)
        n_sheets = len([s for s in entry["sheets"] if s.get("component_type") != "skip"])
        extra = ""
        if entry.get("procedure_item_count") is not None:
            extra = f", 核对程序 {entry['procedure_item_count']} 条"
        print(f"audited {code}: {entry['filename']} ({n_sheets} data sheet(s){extra})")
    if missing:
        print(f"\n[WARN] {len(missing)} file(s) missing (data-blocked): {', '.join(missing)}")
    else:
        print(f"\nAll {len(args.wp_codes)} A17 file(s) audited -> {OUT.relative_to(ROOT)}")


if __name__ == "__main__":
    main()

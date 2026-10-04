# -*- coding: utf-8 -*-
"""I 循环 6 条 entry 真 OnlyOffice 引擎往返验证（HTML → OO → HTML）。

spec: `i-cycle-sync-foundation-and-first-canary` Task 22 · `i1-i3-…` Task 18 · `i2-i4-i5-…` Task 19
范式：照 `verify_l1_oo94_roundtrip.py`（LR-P24），按 entry 参数化。

═══ 链路（每一步都是生产代码，无合成替身）═══

  ① provider `attach_pilot_adapters()` 按 manifest + 真库 published representation 挂 adapter
  ② `CanonicalResolutionService.resolve(extract)` 取当前 representation 的真实字节
  ③ HTML store 载荷 → `build_store_projection` → overlay 发布态
  ④ `adapter.materialize` 写出 xlsx
  ⑤ 🔴 交给真 OnlyOffice 引擎重存（`ConvertService.ashx` xlsx→xlsx；容器经
     `host.docker.internal` 拉本机临时 HTTP 文件）
  ⑥ `adapter.extract` 反读 OO 产物 → G1 等值门 `_assert_roundtrip_equivalent`
  ⑦ `merge_projection_into_store_rows` 合并回 store ⇒ 与输入载荷逐字段比（含行数 = 无幽灵行）
  ⑧ OO 重存后受管行的 formula 列**仍是公式**（formula_mask 生效）；
     非受管 sheet 的公式格 OO 改写数 == 0
  ⑨ 真库 `checklist_responses` 中本 entry 主表键前后快照 digest 不变（本脚本**不写库**）

🔴 载荷是**合成的**（真库 6 个主表键 0 行 —— lane spec 明令标 `synthetic_payload`）。
   合成载荷按各 spec 的 field_specs 生成（text / amount 两种 value_type 全覆盖），
   lane 1 的三条件：≥5 行且行身份互不相同 · 混入 1 条旧格式身份（`cgu-3` 形态）。

用法（仓库根）::

    .venv\\Scripts\\python.exe backend/scripts/e2e/verify_i_cycle_oo94_roundtrip.py [--only i3]

退出码 0 = 全绿；1 = 任一 entry 任一断言失败（打印环节与形态，不降级为合成通过）。
"""
from __future__ import annotations

import argparse
import asyncio
import functools
import hashlib
import http.server
import importlib
import json
import os
import socket
import sys
import tempfile
import threading
import urllib.request
import uuid
from pathlib import Path
from typing import Any

for _stream in (sys.stdout, sys.stderr):
    try:
        _stream.reconfigure(encoding="utf-8")  # type: ignore[attr-defined]
    except Exception:  # pragma: no cover
        pass

_REPO = Path(__file__).resolve().parents[3]
_BACKEND = _REPO / "backend"
if str(_BACKEND) not in sys.path:
    sys.path.insert(0, str(_BACKEND))
os.environ.setdefault("DB_DISABLE_SSL", "True")

import sqlalchemy as sa  # noqa: E402

OO_URL = os.environ.get("I_OO_URL", "http://localhost:8080")
HOST_FROM_CONTAINER = os.environ.get("I_OO_HOST_ALIAS", "host.docker.internal")
N_ROWS = 5

PROVIDERS: dict[str, str] = {
    "i1": "phase5_i1_intangible_assets",
    "i2": "phase5_i2_development_expenditure",
    "i3": "phase5_i3_goodwill",
    "i4": "phase5_i4_long_term_prepaid",
    "i5": "phase5_i5_other_noncurrent_assets",
    "i6": "phase5_i6_research_development_expense",
}


def _free_port() -> int:
    with socket.socket() as sock:
        sock.bind(("0.0.0.0", 0))
        return int(sock.getsockname()[1])


def oo_resave(src: Path, dst: Path) -> dict[str, Any]:
    """把 src 交给真 OO 引擎 xlsx→xlsx 重存，结果写 dst。"""
    port = _free_port()
    handler = functools.partial(http.server.SimpleHTTPRequestHandler, directory=str(src.parent))
    server = http.server.ThreadingHTTPServer(("0.0.0.0", port), handler)
    threading.Thread(target=server.serve_forever, daemon=True).start()
    try:
        body = {
            "async": False, "filetype": "xlsx", "outputtype": "xlsx",
            "key": uuid.uuid4().hex, "title": src.name,
            "url": f"http://{HOST_FROM_CONTAINER}:{port}/{src.name}",
        }
        req = urllib.request.Request(
            f"{OO_URL}/ConvertService.ashx",
            data=json.dumps(body).encode("utf-8"),
            headers={"Content-Type": "application/json", "Accept": "application/json"},
            method="POST",
        )
        with urllib.request.urlopen(req, timeout=180) as resp:
            result = json.loads(resp.read().decode("utf-8"))
        if not result.get("endConvert") or not result.get("fileUrl"):
            raise RuntimeError(f"OO ConvertService 未完成转换: {result}")
        with urllib.request.urlopen(str(result["fileUrl"]), timeout=120) as resp:
            dst.write_bytes(resp.read())
        return result
    finally:
        server.shutdown()


def synth_rows(provider: Any, *, row_ids: list[str] | None = None) -> list[dict[str, Any]]:
    """按 spec 的 field_specs 合成 N 行；多表共用一个 store 时字段并集写进同一行。"""
    def set_json_path(root: dict[str, Any], path: str, value: Any) -> None:
        # 合成载荷专用：生产 `json_path.set_json_path` 要求中间容器已存在（`months/0` 需先有 list）
        tokens = path.split("/")
        cur: Any = root
        for i, tok in enumerate(tokens):
            last = i == len(tokens) - 1
            nxt_is_idx = (not last) and tokens[i + 1].isdigit()
            if isinstance(cur, list):
                k = int(tok)
                while len(cur) <= k:
                    cur.append(None)
                if last:
                    cur[k] = value
                else:
                    if cur[k] is None:
                        cur[k] = [] if nxt_is_idx else {}
                    cur = cur[k]
            else:
                if last:
                    cur[tok] = value
                else:
                    cur = cur.setdefault(tok, [] if nxt_is_idx else {})

    specs = provider.managed_row_table_specs()
    id_key = specs[0].row_identity_key
    rows: list[dict[str, Any]] = []
    for n in range(1, N_ROWS + 1):
        # 🔴 第 3 行用旧格式位置化身份（lane 1 grandfather 条件）：它必须原样往返
        rid = (
            row_ids[n - 1]
            if row_ids is not None
            else ("cgu-3" if n == 3 else f"i-rt-{provider.ADAPTER_ID.split('.')[0]}-{n:02d}-{uuid.UUID(int=n).hex[-6:]}")
        )
        row: dict[str, Any] = {id_key: rid}
        for spec in specs:
            for idx, f in enumerate(spec.field_specs):
                _fid, _col, kind, vtype, path = f[0], f[1], f[2], f[3], f[4]
                if kind != "editable" or not path:
                    continue
                value: Any = f"往返{n}-{_fid}" if vtype == "text" else float(1000 * n + idx * 7)
                set_json_path(row, path, value)
        rows.append(row)
    return rows


def _flatten(row: dict[str, Any], prefix: str = "") -> dict[str, Any]:
    out: dict[str, Any] = {}
    for k, v in row.items():
        key = f"{prefix}{k}"
        if isinstance(v, dict):
            out.update(_flatten(v, key + "/"))
        elif isinstance(v, list):
            for i, item in enumerate(v):
                if isinstance(item, dict):
                    out.update(_flatten(item, f"{key}/{i}/"))
                else:
                    out[f"{key}/{i}"] = item
        else:
            out[key] = v
    return out


def _same(a: Any, b: Any) -> bool:
    if isinstance(a, (int, float)) and b not in (None, ""):
        try:
            return abs(float(a) - float(b)) < 1e-9
        except (TypeError, ValueError):
            return False
    return str(a) == str(b)


def _formula_cells(path: Path) -> dict[str, dict[str, str]]:
    import openpyxl

    wb = openpyxl.load_workbook(path, data_only=False)
    out: dict[str, dict[str, str]] = {}
    for ws in wb.worksheets:
        cells: dict[str, str] = {}
        for row in ws.iter_rows():
            for c in row:
                v = c.value
                text = getattr(v, "text", v)
                if isinstance(text, str) and text.startswith("="):
                    cells[c.coordinate] = text
        out[ws.title] = cells
    wb.close()
    return out


def _managed_rows_formula_report(path: Path, provider: Any, contract: Any) -> list[str]:
    """受管行 formula 列必须仍是公式；契约已登记的 `registered_not_fixed` 缺公式除外。"""
    import openpyxl

    allowed_missing: set[str] = set()
    review = dict(getattr(contract, "canonical_payload", {}).get("review") or {})
    raw_quirks = review.get("known_template_quirks") or []
    quirks = raw_quirks.values() if isinstance(raw_quirks, dict) else raw_quirks
    for quirk in quirks:
        if not isinstance(quirk, dict):
            continue
        if quirk.get("handling") != "registered_not_fixed" or quirk.get("actual") is not None:
            continue
        allowed_missing.update(str(cell) for cell in (quirk.get("cells") or ()))

    wb = openpyxl.load_workbook(path, data_only=False)
    bad: list[str] = []
    for spec in provider.managed_row_table_specs():
        ws = wb[spec.managed_sheet]
        hit = [
            r for r in range(spec.first_data_row, ws.max_row + 1)
            if ws[f"{spec.uuid_col}{r}"].value not in (None, "")
            and r < (spec.footer_row + N_ROWS * 2)
        ]
        hit = hit[:N_ROWS]
        if len(hit) != N_ROWS:
            bad.append(f"{spec.table_key}: 带 UUID 的受管行 {len(hit)} 行（期望 {N_ROWS}）")
            continue
        for r in hit:
            for col in spec.formula_columns:
                v = ws[f"{col}{r}"].value
                text = getattr(v, "text", v)
                if not (isinstance(text, str) and text.startswith("=")):
                    cell_ref = f"{spec.managed_sheet}!{col}{r}"
                    if cell_ref not in allowed_missing:
                        bad.append(f"{cell_ref}={text!r}")
    wb.close()
    return bad


async def _store_digest(session: Any, item_ids: tuple[str, ...]) -> tuple[int, str]:
    rows = (
        await session.execute(
            sa.text(
                "SELECT item_id, wp_id::text, coalesce(remark,''), coalesce(conclusion,'') "
                "FROM checklist_responses WHERE item_id = ANY(:ids) ORDER BY item_id, wp_id"
            ),
            {"ids": list(item_ids)},
        )
    ).fetchall()
    h = hashlib.sha256(json.dumps([list(r) for r in rows], ensure_ascii=False).encode()).hexdigest()
    return len(rows), h


async def run_one(tag: str) -> dict[str, Any]:
    from app.core.database import async_session
    from app.services.workpaper_sync.adapters.registry import WorkpaperSyncAdapterRegistry
    from app.services.workpaper_sync.content_mutation import ContentMutationService
    from app.services.workpaper_sync.entry_profile import load_entry_manifest
    from app.services.workpaper_sync.resolution import (
        CanonicalArtifactRepository,
        CanonicalResolutionService,
        ResolutionIntent,
    )
    from app.services.workpaper_sync.store_projection_response import (
        _overlay_with_published_substrate,
    )

    P = importlib.import_module(f"app.services.workpaper_sync.{PROVIDERS[tag]}")
    rep: dict[str, Any] = {"entry": P.ENTRY_ID, "steps": {}}

    def fail(step: str, msg: str) -> dict[str, Any]:
        rep.update(ok=False, failed_step=step, error=msg)
        print(f"  [{tag} {step}] ❌ {msg}")
        return rep

    async with async_session() as session:
        target = (
            await session.execute(
                sa.text(
                    "SELECT es.wp_id, wp.project_id, es.representation_generation "
                    "FROM working_paper_sync_entry_state es "
                    "JOIN working_paper wp ON wp.id = es.wp_id WHERE es.entry_id = :e"
                ),
                {"e": P.ENTRY_ID},
            )
        ).first()
        if target is None:
            return fail("0", "entry_state 无本 entry 行 —— 先完成首版发布")
        wp_id, project_id = target.wp_id, target.project_id
        item_ids = tuple(P.all_store_item_ids())
        before = await _store_digest(session, item_ids)
        rep["wp_id"], rep["generation"] = str(wp_id), int(target.representation_generation)

        registry = WorkpaperSyncAdapterRegistry(manifest=load_entry_manifest())
        ids = await P.attach_pilot_adapters(registry, session=session)
        if ids != (P.ADAPTER_ID,):
            return fail("①", f"attach 返回 {ids!r}")
        registration = registry.resolve_for_entry(P.ENTRY_ID)
        contract = registration.contract
        rep["steps"]["attach"] = True

        resolver = CanonicalResolutionService(session, CanonicalArtifactRepository(_BACKEND))
        resolution = await resolver.resolve(
            intent=ResolutionIntent.extract, project_id=project_id, wp_id=wp_id, entry_id=P.ENTRY_ID
        )
        substrate = resolution.artifact_path

        # I4 的物理数据区带横向 shared formula（O19:R22），安全门禁止为新身份插行时猜列位移。
        # 取发布 substrate 的既存主表身份做 in-place roundtrip；I1/I2/I3/I5/I6 仍用新身份覆盖插入路径。
        existing_ids: list[str] | None = None
        if tag == "i4":
            base_projection = registration.adapter.extract(artifact=substrate, contract=contract)
            existing_ids = list(base_projection.row_keys.get(P.ROWS_TABLE_KEY, ()))[:N_ROWS]
            if len(existing_ids) != N_ROWS:
                return fail("②", f"I4 substrate 既存身份只有 {len(existing_ids)} 个（需要 {N_ROWS}）")
            rep["steps"]["uses_existing_substrate_identities"] = True
        rows = synth_rows(P, row_ids=existing_ids)
        id_key = P.managed_row_table_specs()[0].row_identity_key
        payload = json.dumps(rows, ensure_ascii=False)
        store_projection = P.build_store_projection(payload, contract=contract)
        projection, _ = await _overlay_with_published_substrate(
            resolution=resolver, project_id=project_id, wp_id=wp_id, entry_id=P.ENTRY_ID,
            registration=registration, store_projection=store_projection,
        )
        n_tables = len(P.managed_row_table_specs())
        n_rows = sum(len(v) for v in projection.row_keys.values())
        if n_rows != N_ROWS * n_tables:
            return fail("③", f"projection 行数 {n_rows} ≠ {N_ROWS}×{n_tables}")
        rep["steps"]["projection"] = {"values": len(projection.values), "rows": n_rows}

        with tempfile.TemporaryDirectory() as tmp:
            mat = Path(tmp) / f"{tag}-materialized.xlsx"
            oo = Path(tmp) / f"{tag}-oo-resaved.xlsx"
            registration.adapter.materialize(
                substrate=substrate, projection=projection, output=mat, contract=contract
            )
            mat_formulas = _formula_cells(mat)
            try:
                oo_result = oo_resave(mat, oo)
            except Exception as exc:  # noqa: BLE001 —— 真栈须如实报告
                return fail("⑤", f"OO ConvertService 失败 {type(exc).__name__}: {exc}")
            rep["steps"]["oo_resave"] = {"size": oo.stat().st_size, "percent": oo_result.get("percent")}

            extracted = registration.adapter.extract(artifact=oo, contract=contract)
            try:
                ContentMutationService._assert_roundtrip_equivalent(
                    None,  # type: ignore[arg-type]
                    intended=projection, extracted=extracted, contract=contract,
                )
            except Exception as exc:  # noqa: BLE001
                return fail("⑥", f"G1 等值门失败 {type(exc).__name__}: {exc}")
            rep["steps"]["g1_equivalent"] = True

            # 框架 merge 可能就地更新 base row dict；保存不可变输入快照，验收不能让被测函数改写 expected。
            expected_rows = json.loads(json.dumps(rows, ensure_ascii=False))
            merged, *_ = P.merge_projection_into_store_rows(projection=extracted, base_rows=rows)
            by_id = {r.get(id_key): r for r in merged}
            diffs: list[str] = []
            for src in expected_rows:
                got = by_id.get(src[id_key])
                if got is None:
                    diffs.append(f"{src[id_key]} 丢失")
                    continue
                fs, fg = _flatten(src), _flatten(got)
                for k, v in fs.items():
                    if not _same(v, fg.get(k)):
                        diffs.append(f"{src[id_key][-6:]}.{k}: {v!r} → {fg.get(k)!r}")
            if diffs:
                return fail("⑦", f"store 往返不等 {len(diffs)} 处: {diffs[:6]}")
            # 模板预印业务行（如 I6 的人工费/材料费等）不是 ghost：它们有 GTROW-* 模板身份 +
            # 非空业务锚点，materialize 默认 clear 语义应保留。真正的 ghost 是锚点全空的模板骨架行。
            # 因此判据是「输入 N 行逐字段全在且不变」+「extra 只能是非空 template identity」。
            src_ids = {r[id_key] for r in expected_rows}
            extras = [r for r in merged if r.get(id_key) not in src_ids]
            specs = P.managed_row_table_specs()
            anchor_paths = [s.field_specs[s.ghost_row_anchor_index][4] for s in specs]
            flat_extras = [_flatten(r) for r in extras]
            invalid_extras = [
                r for r, flat in zip(extras, flat_extras)
                if not str(r.get(id_key) or "").startswith("GTROW-")
                or not any(str(flat.get(path) or "").strip() for path in anchor_paths)
            ]
            if invalid_extras:
                return fail(
                    "⑦",
                    f"合并回 store 出现 {len(invalid_extras)} 个空骨架/非模板 extra："
                    f"{json.dumps([_flatten(r) for r in invalid_extras[:4]], ensure_ascii=False)[:600]}",
                )
            if extras:
                rep["steps"]["preserved_template_business_rows"] = len(extras)
            if tag in {"i1", "i3"} and "cgu-3" not in by_id:
                return fail("⑦", "旧格式身份 cgu-3 未原样往返（grandfather 失效）")
            rep["steps"]["store_equal"] = {"rows": N_ROWS, "keys_per_row": len(_flatten(expected_rows[0]))}

            bad = _managed_rows_formula_report(oo, P, contract)
            if bad:
                return fail("⑧a", f"受管行 formula 列非公式 {len(bad)} 处: {bad[:6]}")
            oo_formulas = _formula_cells(oo)
            managed = {s.managed_sheet for s in P.managed_row_table_specs()}
            changed: list[str] = []
            for sheet, cells in mat_formulas.items():
                if sheet in managed:
                    continue
                after = oo_formulas.get(sheet, {})
                for coord, f in cells.items():
                    if after.get(coord) != f:
                        changed.append(f"{sheet}!{coord}")
            rep["steps"]["unmanaged_formula_cells"] = sum(
                len(v) for k, v in mat_formulas.items() if k not in managed
            )
            rep["steps"]["unmanaged_formula_changed_by_oo"] = len(changed)
            if changed:
                return fail("⑧b", f"OO 改写了非受管 sheet 公式 {len(changed)} 格: {changed[:6]}")

        after = await _store_digest(session, item_ids)
        if after != before:
            return fail("⑨", f"主表键快照变化 {before} → {after}")
        rep["steps"]["db_untouched"] = {"rows": after[0]}
    rep["ok"] = True
    rep["sync_test_run_id"] = f"i-oo94-{tag}-{uuid.uuid4().hex[:12]}"
    print(f"  [{tag}] ✅ {P.ENTRY_ID} {json.dumps(rep['steps'], ensure_ascii=False)}")
    return rep


async def main_async(only: str | None, out: Path | None) -> int:
    reports = []
    for tag in PROVIDERS:
        if only and tag != only:
            continue
        reports.append(await run_one(tag))
    if out:
        out.write_text(json.dumps({
            "evidence_kind": "synthetic_payload_no_live_db_baseline",
            "oo_url": OO_URL,
            "engine": "OnlyOffice ConvertService xlsx->xlsx (docker audit-onlyoffice)",
            "reports": reports,
        }, ensure_ascii=False, indent=2), encoding="utf-8")
    ok = all(r.get("ok") for r in reports)
    print("✅ I 循环真 OO 引擎往返全绿" if ok else "❌ 有 entry 未通过")
    return 0 if ok else 1


if __name__ == "__main__":
    ap = argparse.ArgumentParser()
    ap.add_argument("--only", default=None)
    ap.add_argument("--json", default=None)
    a = ap.parse_args()
    raise SystemExit(asyncio.run(main_async(a.only, Path(a.json) if a.json else None)))

"""公式管理浏览器门禁的 fresh-DB 确定性 seed。

职责严格限定为：可登录 admin、一个 standalone 项目、A1 与 D2-1 两个有效 XLSX
底稿。合并项目由 ``scripts/seed/seed_consol_uat.py`` 独立负责。
"""

from __future__ import annotations

import argparse
import asyncio
import json
import sys
from datetime import date
from pathlib import Path

# 允许以脚本路径直接执行，把 backend/ 加入模块搜索路径。
ROOT = Path(__file__).resolve().parents[2]
if str(ROOT) not in sys.path:
    sys.path.insert(0, str(ROOT))

import sqlalchemy as sa
from openpyxl import Workbook, load_workbook
from sqlalchemy.ext.asyncio import AsyncSession, async_sessionmaker, create_async_engine

from app.core.config import settings
from app.core.security import hash_password, verify_password
from app.models.base import UserRole
from app.models.core import Project, ProjectStatus, ProjectType, User
from app.models.extension_models import AccountingStandard  # noqa: F401 — 注册 Project FK metadata
from app.models.workpaper_models import WorkingPaper, WpIndex, WpSourceType, WpStatus

ADMIN_USERNAME = "admin"
ADMIN_PASSWORD = "admin123"
ADMIN_EMAIL = "admin@gt.com"
PROJECT_NAME_PREFIX = "【E2E】公式管理运行时门禁"
WP_SPECS = {
    "A1": "财务报表审计总览",
    "D2-1": "应收账款审定表",
}


def _database_url() -> str:
    url = str(settings.DATABASE_URL)
    return url.replace("postgresql://", "postgresql+asyncpg://", 1) if url.startswith("postgresql://") else url


def _project_name(year: int) -> str:
    return f"{PROJECT_NAME_PREFIX}_{year}"


def _ensure_workbook(path: Path, code: str) -> None:
    """写入最小合法 XLSX；已有合法文件保持不变，seed 可幂等重放。"""
    if path.is_file() and path.stat().st_size > 0:
        try:
            book = load_workbook(path, read_only=True, data_only=False)
            book.close()
            return
        except Exception:
            pass
    path.parent.mkdir(parents=True, exist_ok=True)
    book = Workbook()
    sheet = book.active
    sheet.title = code
    sheet["A1"] = "公式管理 E2E 门禁"
    sheet["A2"] = code
    book.save(path)
    book.close()


async def _ensure_admin(db: AsyncSession, stats: dict[str, list[str]]) -> User:
    admin = (await db.execute(sa.select(User).where(User.username == ADMIN_USERNAME))).scalars().first()
    if admin is None:
        email_owner = (await db.execute(sa.select(User).where(User.email == ADMIN_EMAIL))).scalars().first()
        if email_owner is not None:
            raise RuntimeError(f"邮箱 {ADMIN_EMAIL} 已被非 admin 用户占用")
        admin = User(
            username=ADMIN_USERNAME,
            email=ADMIN_EMAIL,
            hashed_password=hash_password(ADMIN_PASSWORD),
            role=UserRole.admin,
            is_active=True,
            is_deleted=False,
        )
        db.add(admin)
        await db.flush()
        stats["created"].append("user:admin")
    else:
        changed: list[str] = []
        if not verify_password(ADMIN_PASSWORD, admin.hashed_password):
            admin.hashed_password = hash_password(ADMIN_PASSWORD)
            changed.append("credentials")
        if admin.role != UserRole.admin:
            admin.role = UserRole.admin
            changed.append("role")
        if not admin.is_active or admin.is_deleted:
            admin.is_active = True
            admin.is_deleted = False
            changed.append("active")
        target = stats["updated"] if changed else stats["skipped"]
        suffix = f"({','.join(changed)})" if changed else ""
        target.append(f"user:admin{suffix}")
    return admin


async def _ensure_project(
    db: AsyncSession,
    *,
    year: int,
    admin: User,
    stats: dict[str, list[str]],
) -> Project:
    name = _project_name(year)
    rows = list((await db.execute(sa.select(Project).where(
        Project.client_name == name,
        Project.audit_year == year,
        Project.is_deleted == sa.false(),
    ))).scalars().all())
    if len(rows) > 1:
        raise RuntimeError(f"公式管理 E2E 项目重复：{name} 共 {len(rows)} 条")
    if rows:
        project = rows[0]
        stats["skipped"].append(f"project:{name}")
    else:
        project = Project(
            name=name,
            client_name=name,
            audit_period_start=date(year, 1, 1),
            audit_period_end=date(year, 12, 31),
            audit_year=year,
            project_type=ProjectType.annual,
            status=ProjectStatus.planning,
            manager_id=admin.id,
            partner_id=admin.id,
            business_category="A1",
            scenario="normal",
            template_type="soe",
            report_scope="standalone",
            applicable_standard_v2={"entity_type": "soe", "scope": "standalone", "stage": "normal"},
            is_deleted=False,
        )
        db.add(project)
        await db.flush()
        stats["created"].append(f"project:{name}")

    project.manager_id = admin.id
    project.partner_id = admin.id
    project.audit_period_start = project.audit_period_start or date(year, 1, 1)
    project.audit_period_end = project.audit_period_end or date(year, 12, 31)
    project.audit_year = year
    project.report_scope = "standalone"
    project.template_type = "soe"
    project.is_deleted = False
    return project


async def _ensure_papers(
    db: AsyncSession,
    *,
    project: Project,
    admin: User,
    year: int,
    stats: dict[str, list[str]],
) -> dict[str, str]:
    ids: dict[str, str] = {}
    storage_root = Path(__file__).resolve().parents[2] / "storage" / "e2e" / "formula-manager" / str(project.id)

    for code, label in WP_SPECS.items():
        index = (await db.execute(sa.select(WpIndex).where(
            WpIndex.project_id == project.id,
            WpIndex.wp_code == code,
        ))).scalars().first()
        if index is None:
            index = WpIndex(
                project_id=project.id,
                wp_code=code,
                wp_name=label,
                audit_cycle=code[0],
                assigned_to=admin.id,
                reviewer=admin.id,
                status=WpStatus.not_started,
                is_deleted=False,
            )
            db.add(index)
            await db.flush()
            stats["created"].append(f"wp_index:{code}")
        else:
            index.is_deleted = False
            stats["skipped"].append(f"wp_index:{code}")

        path = storage_root / f"{code}.xlsx"
        _ensure_workbook(path, code)
        paper = (await db.execute(sa.select(WorkingPaper).where(
            WorkingPaper.project_id == project.id,
            WorkingPaper.wp_index_id == index.id,
        ))).scalars().first()
        if paper is None:
            paper = WorkingPaper(
                project_id=project.id,
                wp_index_id=index.id,
                file_path=str(path),
                source_type=WpSourceType.template,
                file_version=1,
                assigned_to=admin.id,
                reviewer=admin.id,
                created_by=admin.id,
                updated_by=admin.id,
                parsed_data={"html_data": {code: {"cells": {}}}},
                is_deleted=False,
            )
            db.add(paper)
            await db.flush()
            stats["created"].append(f"working_paper:{code}")
        else:
            paper.file_path = str(path)
            paper.is_deleted = False
            paper.assigned_to = admin.id
            stats["skipped"].append(f"working_paper:{code}")
        ids[code] = str(paper.id)

    await db.flush()
    return ids


async def seed(year: int) -> dict:
    engine = create_async_engine(_database_url())
    Session = async_sessionmaker(engine, expire_on_commit=False)
    stats: dict[str, list[str]] = {"created": [], "updated": [], "skipped": []}
    try:
        async with Session() as db:
            admin = await _ensure_admin(db, stats)
            project = await _ensure_project(db, year=year, admin=admin, stats=stats)
            paper_ids = await _ensure_papers(db, project=project, admin=admin, year=year, stats=stats)
            await db.commit()

        async with Session() as verify:
            project = (await verify.execute(sa.select(Project).where(
                Project.client_name == _project_name(year),
                Project.audit_year == year,
                Project.report_scope == "standalone",
                Project.is_deleted == sa.false(),
            ))).scalars().one()
            rows = (await verify.execute(
                sa.select(WpIndex.wp_code, WorkingPaper.id, WorkingPaper.file_path)
                .join(WorkingPaper, WorkingPaper.wp_index_id == WpIndex.id)
                .where(
                    WpIndex.project_id == project.id,
                    WpIndex.is_deleted == sa.false(),
                    WorkingPaper.is_deleted == sa.false(),
                    WpIndex.wp_code.in_(list(WP_SPECS)),
                )
            )).all()
            by_code = {code: (paper_id, file_path) for code, paper_id, file_path in rows}
            if set(by_code) != set(WP_SPECS):
                raise RuntimeError(f"底稿 seed 不完整：expected={sorted(WP_SPECS)}, actual={sorted(by_code)}")
            for code, (_, file_path) in by_code.items():
                _ensure_workbook(Path(file_path), code)
            admin = (await verify.execute(sa.select(User).where(
                User.username == ADMIN_USERNAME,
                User.role == UserRole.admin,
                User.is_active == sa.true(),
                User.is_deleted == sa.false(),
            ))).scalars().one()
            return {
                "year": year,
                "project_id": str(project.id),
                "project_name": project.name,
                "admin_id": str(admin.id),
                "wp_ids": {code: str(value[0]) for code, value in by_code.items()},
                "stats": stats,
            }
    finally:
        await engine.dispose()


def main() -> int:
    parser = argparse.ArgumentParser(description="公式管理 runtime browser gate fresh-DB seed")
    parser.add_argument("--year", type=int, default=2025)
    args = parser.parse_args()
    result = asyncio.run(seed(args.year))
    print(json.dumps(result, ensure_ascii=False, sort_keys=True, indent=2))
    return 0


if __name__ == "__main__":
    sys.exit(main())

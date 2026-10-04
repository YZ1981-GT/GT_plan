# -*- coding: utf-8 -*-
"""L7 其他非流动负债（受管 明细表L7-2）真 OnlyOffice 引擎往返验证（HTML → OO → HTML）。

spec: l7-true-bidirectional-2026-10-01 · 第6步（复用 L1/L4 的 Task 10 / LR-P24 同一链路）

链路与 `verify_l4_oo94_roundtrip.py` 完全相同（L7-2 有公式列 E/L/M/N/O，用 L4 的「公式幸存」范式，
不照抄 L3 的「零公式列」断言）：生产 attach（断言 ids==('l7.other_noncurrent_liabilities',)）→
真 representation substrate → HTML 载荷（2 行，稳定 rowId，_row 跳过 formula 列、置 itemName 等
editable 列）→ materialize → docker `audit-onlyoffice` 的 `ConvertService.ashx` xlsx→xlsx 重存 →
extract → G1 等值门 → merge 回 store 逐字段比（rowId 稳定、无幽灵行、行命中数==输入数）→
受管行 E/L/M/N/O 五列公式往返后仍是公式 → 不写库（前后 L7-% 行数不变）。9 步。

退出码 0 = 全绿；1 = 任一断言失败（打印失败环节，不降级为合成通过）。
"""
from __future__ import annotations

import asyncio
import importlib.util
import json
import os
import sys
import tempfile
from pathlib import Path
from typing import Any

for _stream in (sys.stdout, sys.stderr):
    try:
        _stream.reconfigure(encoding="utf-8")  # type: ignore[attr-defined]
    except Exception:  # pragma: no cover
        pass

_HERE = Path(__file__).resolve().parent
_BACKEND = _HERE.parents[1]
if str(_BACKEND) not in sys.path:
    sys.path.insert(0, str(_BACKEND))
os.environ.setdefault("DB_DISABLE_SSL", "True")

_spec = importlib.util.spec_from_file_location("_verify_l1", _HERE / "verify_l1_oo94_roundtrip.py")
assert _spec and _spec.loader
_V1 = importlib.util.module_from_spec(_spec)
sys.modules[_spec.name] = _V1
_spec.loader.exec_module(_V1)

import sqlalchemy as sa  # noqa: E402

ENTRY_ID = "xlsx/gt-l7-other-noncurrent-liabilities"


def _row(rid: str, name: str, base: dict[str, Any]) -> dict[str, Any]:
    from app.services.workpaper_sync.phase5_l7_sheets import MANAGED_FIELD_SPECS_7

    row: dict[str, Any] = {"rowId": rid}
    for _key, _col, mode, vtype, jkey, *_ in MANAGED_FIELD_SPECS_7:
        if mode == "formula":
            continue
        row[jkey] = base.get(jkey, 0 if vtype in ("amount", "rate", "integer") else "")
    row["itemName"] = name
    return row


ROWS: list[dict[str, Any]] = [
    _row("7f0c1e8a-1a2b-4c3d-8e9f-0000000000b1", "递延收益（政府补助）", {
        "beginUnadjusted": 800000, "beginAje": 0, "beginRje": 0,
        "ajeIncrease": 0, "rjeIncrease": 50000,
        "endAjeIncrease": 1200, "endRjeDecrease": 0, "endAjeDecrease": 0, "endRjeIncrease": 0,
        "nature": "政府补助递延", "reason": "财政专项补助分期确认", "maturityInfo": "按项目周期摊销",
    }),
    _row("7f0c1e8a-1a2b-4c3d-8e9f-0000000000b2", "长期应付款（保证金）", {
        "beginUnadjusted": 300000, "beginAje": 0, "beginRje": 0,
        "ajeIncrease": 0, "rjeIncrease": 0,
        "endAjeIncrease": 0, "endRjeDecrease": 20000, "endAjeDecrease": 0, "endRjeIncrease": 0,
        "nature": "履约保证金", "reason": "供应商履约保证", "maturityInfo": "合同到期退还",
    }),
]


async def run() -> int:
    from app.core.database import async_session
    from app.services.workpaper_sync import phase5_l7_other_noncurrent_liabilities as L7
    from app.services.workpaper_sync.adapters.registry import WorkpaperSyncAdapterRegistry
    from app.services.workpaper_sync.content_mutation import ContentMutationService
    from app.services.workpaper_sync.entry_profile import load_entry_manifest
    from app.services.workpaper_sync.resolution import (
        CanonicalArtifactRepository,
        CanonicalResolutionService,
        ResolutionIntent,
    )
    from app.services.workpaper_sync.store_projection_response import (
        _overlay_with_published_substrate,
    )

    fail = _V1._fail
    async with async_session() as session:
        target = (await session.execute(sa.text(
            "SELECT es.wp_id, wp.project_id FROM working_paper_sync_entry_state es "
            "JOIN working_paper wp ON wp.id = es.wp_id WHERE es.entry_id = :e"), {"e": ENTRY_ID})).first()
        if target is None:
            return fail("0", "entry_state 无 L7 行 —— 先完成首版发布")
        wp_id, project_id = target.wp_id, target.project_id
        before_rows = (await session.execute(sa.text(
            "SELECT count(*) FROM checklist_responses WHERE item_id LIKE 'L7-%'"))).scalar_one()
        print(f"[0] wp={str(wp_id)[:8]} project={str(project_id)[:8]} L7 载荷快照={before_rows} 行")

        registry = WorkpaperSyncAdapterRegistry(manifest=load_entry_manifest())
        ids = await L7.attach_adapters(registry, session=session)
        if ids != (L7.ADAPTER_ID,):
            return fail("①", f"attach 返回 {ids!r}（manifest capability 或 representation 未就位）")
        registration = registry.resolve_for_entry(ENTRY_ID)
        contract = registration.contract
        print(f"[①] attach={ids}")

        resolver = CanonicalResolutionService(session, CanonicalArtifactRepository(_BACKEND))
        resolution = await resolver.resolve(
            intent=ResolutionIntent.extract, project_id=project_id, wp_id=wp_id, entry_id=ENTRY_ID)
        substrate = resolution.artifact_path
        print(f"[②] generation={resolution.representation_generation} substrate={substrate.name}")

        store_projection = L7.build_store_projection(json.dumps(ROWS, ensure_ascii=False), contract=contract)
        projection, _ = await _overlay_with_published_substrate(
            resolution=resolver, project_id=project_id, wp_id=wp_id, entry_id=ENTRY_ID,
            registration=registration, store_projection=store_projection)
        n_rows = sum(len(v) for v in projection.row_keys.values())
        print(f"[③] projection values={len(projection.values)} rows={n_rows}")
        if n_rows != len(ROWS):
            return fail("③", f"projection 行数 {n_rows} ≠ 输入 {len(ROWS)}")

        with tempfile.TemporaryDirectory() as tmp:
            mat, oo = Path(tmp) / "l7-materialized.xlsx", Path(tmp) / "l7-oo-resaved.xlsx"
            registration.adapter.materialize(substrate=substrate, projection=projection, output=mat,
                                             contract=contract)
            print(f"[④] materialize size={mat.stat().st_size}")
            try:
                res = _V1.oo_resave(mat, oo)
            except Exception as exc:  # noqa: BLE001
                return fail("⑤", f"OO ConvertService 失败 {type(exc).__name__}: {exc}")
            print(f"[⑤] OO resave OK size={oo.stat().st_size} percent={res.get('percent')}")

            extracted = registration.adapter.extract(artifact=oo, contract=contract)
            try:
                ContentMutationService._assert_roundtrip_equivalent(
                    None, intended=projection, extracted=extracted, contract=contract)  # type: ignore[arg-type]
            except Exception as exc:  # noqa: BLE001
                return fail("⑥", f"G1 等值门失败 {type(exc).__name__}: {exc}")
            print(f"[⑥] extract values={len(extracted.values)} · G1 等值门 OK")

            merged, *_ = L7.merge_projection_into_store_rows(projection=extracted, base_rows=ROWS)
            by_id = {r["rowId"]: r for r in merged}
            diffs: list[str] = []
            for src in ROWS:
                got = by_id.get(src["rowId"])
                if got is None:
                    diffs.append(f"rowId {src['rowId'][-4:]} 丢失")
                    continue
                for k, v in src.items():
                    gv = got.get(k)
                    if isinstance(v, (int, float)) and gv not in (None, ""):
                        same = float(gv) == float(v)
                    else:
                        same = str(gv if gv is not None else "") == str(v)
                    if not same:
                        diffs.append(f"{src['rowId'][-4:]}.{k}: {v!r} → {gv!r}")
            if diffs or len(merged) != len(ROWS):
                return fail("⑦", f"merged={len(merged)} 行，不等 {len(diffs)} 处: {diffs[:8]}")
            print(f"[⑦] L7-L7-2-full-data 往返逐字段相等（{len(ROWS)} 行 × {len(ROWS[0])} 键，rowId 稳定，无幽灵行）")

            import openpyxl

            ws = openpyxl.load_workbook(oo, data_only=False)[L7.MANAGED_SHEET]
            names = {r["itemName"] for r in ROWS}
            hit = [r for r in range(1, ws.max_row + 1) if ws[f"A{r}"].value in names]
            bad = {f"{c}{r}": ws[f"{c}{r}"].value for r in hit for c in L7.FORMULA_COLUMNS
                   if not (isinstance(ws[f"{c}{r}"].value, str) and str(ws[f"{c}{r}"].value).startswith("="))}
            if len(hit) != len(ROWS) or bad:
                return fail("⑧", f"受管行={hit} 非公式格={bad}")
            print(f"[⑧] OO 重存后受管行 {hit} 的五列公式 {len(hit) * len(L7.FORMULA_COLUMNS)} 格全部仍是公式")

        after_rows = (await session.execute(sa.text(
            "SELECT count(*) FROM checklist_responses WHERE item_id LIKE 'L7-%'"))).scalar_one()
        if after_rows != before_rows:
            return fail("⑨", f"L7 载荷行数变化 {before_rows} → {after_rows}")
        print(f"[⑨] L7 载荷 {after_rows} 行不变（本脚本不写库）")

    print("✅ L7 真 OO 引擎往返全绿（HTML→OO→HTML）")
    return 0


if __name__ == "__main__":
    raise SystemExit(asyncio.run(run()))

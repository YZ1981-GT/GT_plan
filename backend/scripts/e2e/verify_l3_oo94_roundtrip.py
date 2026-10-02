# -*- coding: utf-8 -*-
"""L3 长期借款（受管 L3-9）真 OnlyOffice 引擎往返验证（HTML → OO → HTML）。

spec: l-cycle-true-adapter-registration · Task 12（对 L3 复用 Task 10 / LR-P24 的同一链路）

链路与 `verify_l4_oo94_roundtrip.py` 相同（真 OO 重存助手 `verify_l1_oo94_roundtrip.oo_resave`
直接复用，不复制）：生产 attach → 真 representation substrate → HTML 载荷（2 行，稳定 rowId）→
materialize → docker `audit-onlyoffice` 的 `ConvertService.ashx` xlsx→xlsx 重存 → extract →
G1 等值门 → merge 回 store 逐字段比 → 不写库。

🔴 与 L4 的唯一差异：L3-9 受管区**零公式列**（`FORMULA_COLUMNS == ()`）⇒ 第 ⑧ 步不是「公式
幸存」而是「受管行无任何公式格残留」+「物理行命中数 == 输入行数」。照抄 L4 的公式幸存断言会
在空 FORMULA_COLUMNS 上恒真（假绿）。

退出码 0 = 全绿；1 = 任一断言失败（打印失败环节，不降级为合成通过）。
"""
from __future__ import annotations

import asyncio
import importlib.util
import json
import os
import sys
import tempfile
from pathlib import Path
from typing import Any

for _stream in (sys.stdout, sys.stderr):
    try:
        _stream.reconfigure(encoding="utf-8")  # type: ignore[attr-defined]
    except Exception:  # pragma: no cover
        pass

_HERE = Path(__file__).resolve().parent
_BACKEND = _HERE.parents[1]
if str(_BACKEND) not in sys.path:
    sys.path.insert(0, str(_BACKEND))
os.environ.setdefault("DB_DISABLE_SSL", "True")

_spec = importlib.util.spec_from_file_location("_verify_l1", _HERE / "verify_l1_oo94_roundtrip.py")
assert _spec and _spec.loader
_V1 = importlib.util.module_from_spec(_spec)
sys.modules[_spec.name] = _V1
_spec.loader.exec_module(_V1)

import sqlalchemy as sa  # noqa: E402

ENTRY_ID = "xlsx/gt-l3-long-term-loans"


def _row(rid: str, voucher: str, base: dict[str, Any]) -> dict[str, Any]:
    from app.services.workpaper_sync.phase5_l3_sheets import MANAGED_FIELD_SPECS_7

    row: dict[str, Any] = {"rowId": rid}
    for _key, _col, mode, vtype, jkey, *_ in MANAGED_FIELD_SPECS_7:
        if mode == "formula":
            continue
        if vtype == "boolean":
            row[jkey] = base.get(jkey, False)
        elif vtype == "amount":
            row[jkey] = base.get(jkey, 0)
        else:
            row[jkey] = base.get(jkey, "")
    row["voucherNo"] = voucher
    return row


ROWS: list[dict[str, Any]] = [
    _row("7a1c2e9b-3d4e-4f50-8a1b-0000000000c1", "记-2025-0312", {
        "date": "2025-03-12", "businessContent": "取得三年期项目贷款", "counterAccount": "银行存款",
        "counterSubAccount": "工行基本户", "debitAmount": 0, "creditAmount": 2000000,
        "supportingDoc": "借款合同+银行回单", "check1": True, "check2": True, "check3": True,
        "check4": True, "check5": True, "indexNo": "L3-9-01", "abnormal": False, "remark": "往返验证行一",
    }),
    _row("7a1c2e9b-3d4e-4f50-8a1b-0000000000c2", "记-2025-0930", {
        "date": "2025-09-30", "businessContent": "归还长期借款本金", "counterAccount": "银行存款",
        "counterSubAccount": "建行一般户", "debitAmount": 500000, "creditAmount": 0,
        "supportingDoc": "还款申请+银行回单", "check1": True, "check2": True, "check3": False,
        "check4": True, "check5": True, "indexNo": "L3-9-02", "abnormal": True, "remark": "往返验证行二（含异常）",
    }),
]


async def run() -> int:
    from app.core.database import async_session
    from app.services.workpaper_sync import phase5_l3_long_term_loans as L3
    from app.services.workpaper_sync.adapters.registry import WorkpaperSyncAdapterRegistry
    from app.services.workpaper_sync.content_mutation import ContentMutationService
    from app.services.workpaper_sync.entry_profile import load_entry_manifest
    from app.services.workpaper_sync.resolution import (
        CanonicalArtifactRepository,
        CanonicalResolutionService,
        ResolutionIntent,
    )
    from app.services.workpaper_sync.store_projection_response import (
        _overlay_with_published_substrate,
    )

    fail = _V1._fail
    async with async_session() as session:
        target = (await session.execute(sa.text(
            "SELECT es.wp_id, wp.project_id FROM working_paper_sync_entry_state es "
            "JOIN working_paper wp ON wp.id = es.wp_id WHERE es.entry_id = :e"), {"e": ENTRY_ID})).first()
        if target is None:
            return fail("0", "entry_state 无 L3 行 —— 先完成首版发布")
        wp_id, project_id = target.wp_id, target.project_id
        before_rows = (await session.execute(sa.text(
            "SELECT count(*) FROM checklist_responses WHERE item_id LIKE 'L3-%'"))).scalar_one()
        print(f"[0] wp={str(wp_id)[:8]} project={str(project_id)[:8]} L3 载荷快照={before_rows} 行")

        registry = WorkpaperSyncAdapterRegistry(manifest=load_entry_manifest())
        ids = await L3.attach_adapters(registry, session=session)
        if ids != (L3.ADAPTER_ID,):
            return fail("①", f"attach 返回 {ids!r}（manifest capability 或 representation 未就位）")
        registration = registry.resolve_for_entry(ENTRY_ID)
        contract = registration.contract
        print(f"[①] attach={ids}")

        resolver = CanonicalResolutionService(session, CanonicalArtifactRepository(_BACKEND))
        resolution = await resolver.resolve(
            intent=ResolutionIntent.extract, project_id=project_id, wp_id=wp_id, entry_id=ENTRY_ID)
        substrate = resolution.artifact_path
        print(f"[②] generation={resolution.representation_generation} substrate={substrate.name}")

        store_projection = L3.build_store_projection(json.dumps(ROWS, ensure_ascii=False), contract=contract)
        projection, _ = await _overlay_with_published_substrate(
            resolution=resolver, project_id=project_id, wp_id=wp_id, entry_id=ENTRY_ID,
            registration=registration, store_projection=store_projection)
        n_rows = sum(len(v) for v in projection.row_keys.values())
        print(f"[③] projection values={len(projection.values)} rows={n_rows}")
        if n_rows != len(ROWS):
            return fail("③", f"projection 行数 {n_rows} ≠ 输入 {len(ROWS)}")

        with tempfile.TemporaryDirectory() as tmp:
            mat, oo = Path(tmp) / "l3-materialized.xlsx", Path(tmp) / "l3-oo-resaved.xlsx"
            registration.adapter.materialize(substrate=substrate, projection=projection, output=mat,
                                             contract=contract)
            print(f"[④] materialize size={mat.stat().st_size}")
            try:
                res = _V1.oo_resave(mat, oo)
            except Exception as exc:  # noqa: BLE001
                return fail("⑤", f"OO ConvertService 失败 {type(exc).__name__}: {exc}")
            print(f"[⑤] OO resave OK size={oo.stat().st_size} percent={res.get('percent')}")

            extracted = registration.adapter.extract(artifact=oo, contract=contract)
            try:
                ContentMutationService._assert_roundtrip_equivalent(
                    None, intended=projection, extracted=extracted, contract=contract)  # type: ignore[arg-type]
            except Exception as exc:  # noqa: BLE001
                return fail("⑥", f"G1 等值门失败 {type(exc).__name__}: {exc}")
            print(f"[⑥] extract values={len(extracted.values)} · G1 等值门 OK")

            merged, *_ = L3.merge_projection_into_store_rows(projection=extracted, base_rows=ROWS)
            by_id = {r["rowId"]: r for r in merged}
            diffs: list[str] = []
            for src in ROWS:
                got = by_id.get(src["rowId"])
                if got is None:
                    diffs.append(f"rowId {src['rowId'][-4:]} 丢失")
                    continue
                for k, v in src.items():
                    gv = got.get(k)
                    if isinstance(v, bool):
                        same = bool(gv) == v
                    elif isinstance(v, (int, float)) and gv not in (None, ""):
                        same = float(gv) == float(v)
                    else:
                        same = str(gv if gv is not None else "") == str(v)
                    if not same:
                        diffs.append(f"{src['rowId'][-4:]}.{k}: {v!r} → {gv!r}")
            if diffs or len(merged) != len(ROWS):
                return fail("⑦", f"merged={len(merged)} 行，不等 {len(diffs)} 处: {diffs[:8]}")
            print(f"[⑦] L3-L3-9-voucher-rows 往返逐字段相等（{len(ROWS)} 行 × {len(ROWS[0])} 键，rowId 稳定，无幽灵行）")

            # ⑧ L3-9 受管区零公式列 ⇒ 断言「受管行物理命中 == 输入行数」且「A..P 无公式格残留」。
            #    照抄 L4 的公式幸存断言会在空 FORMULA_COLUMNS 上恒真（假绿）。
            import openpyxl
            from openpyxl.utils import get_column_letter

            assert L3.FORMULA_COLUMNS == (), "L3-9 应为零公式列；若模板变更须改本断言"
            ws = openpyxl.load_workbook(oo, data_only=False)[L3.MANAGED_SHEET]
            vouchers = {r["voucherNo"] for r in ROWS}
            hit = [r for r in range(1, ws.max_row + 1) if ws[f"B{r}"].value in vouchers]
            stray_formula = {
                f"{get_column_letter(c)}{r}": ws.cell(r, c).value
                for r in hit for c in range(1, 17)
                if isinstance(ws.cell(r, c).value, str) and str(ws.cell(r, c).value).startswith("=")
            }
            if len(hit) != len(ROWS) or stray_formula:
                return fail("⑧", f"受管行={hit} 残留公式格={stray_formula}")
            print(f"[⑧] OO 重存后受管行 {hit} 的 A..P 共 {len(hit) * 16} 格无公式残留（L3-9 零公式列如期）")

        after_rows = (await session.execute(sa.text(
            "SELECT count(*) FROM checklist_responses WHERE item_id LIKE 'L3-%'"))).scalar_one()
        if after_rows != before_rows:
            return fail("⑨", f"L3 载荷行数变化 {before_rows} → {after_rows}")
        print(f"[⑨] L3 载荷 {after_rows} 行不变（本脚本不写库）")

    print("✅ L3 真 OO 引擎往返全绿（HTML→OO→HTML）")
    return 0


if __name__ == "__main__":
    raise SystemExit(asyncio.run(run()))

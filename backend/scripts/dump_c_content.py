"""Dump text of representative C-category sheets (read-only)."""
from __future__ import annotations

import sys
from pathlib import Path

import openpyxl

BASE = Path(
    r"d:\GT_plan\数据\致同通用审计程序及底稿模板（2025年修订）"
    r"\1.致同审计程序及底稿模板（2025年）\3.风险应对-一般性程序与控制测试（C1-C26）"
)

TARGETS = [
    ("C1 企业层面控制测试.xlsx", ["C1 企业层面控制测试程序表", "C1-4-4企业层面内控测试示例4"]),
    ("C22 IT一般控制测试.xlsx", ["C22 IT一般控制测试", "SA-3", "PE-3a"]),
    ("C23 会计分录 - 控制测试.xlsx", ["C23A 会计分录控制测试程序表", "会计分录控制测试C23-2"]),
    ("C24 会计分录 - 细节测试.xlsx", ["C24-0汇总表", "C24-3完整性-跳号测试", "C24-5细节测试-异常分录测试"]),
    ("C25 利用内审工作.xlsx", ["C25利用内部审计工作"]),
    ("C26 信息处理控制测试.xlsx", ["C26 信息处理控制测试"]),
    (r"C2 销售业务循环控制测试\C2 销售循环控制测试.xlsx", ["C2控制测试汇总表", "C2-1-X控制测试"]),
    (r"C2 销售业务循环控制测试\C2-2 销售循环评价控制偏差.xlsx", ["C2-2评价控制偏差"]),
]


def dump(ws, max_rows: int = 30) -> None:
    for r_i, row in enumerate(ws.iter_rows(values_only=True), start=1):
        if r_i > max_rows:
            print("    ... (truncated)")
            break
        cells = [str(c).strip().replace("\n", " ") for c in row if c is not None and str(c).strip()]
        if cells:
            print(f"    {r_i:>3}: {' | '.join(cells)[:200]}")


def main() -> None:
    for fname, sheets in TARGETS:
        p = BASE / fname
        try:
            wb = openpyxl.load_workbook(p, data_only=False)
        except Exception as e:  # noqa: BLE001
            print(f"### {fname} <ERROR {e}>")
            continue
        for name in sheets:
            if name not in wb.sheetnames:
                print(f"\n### {fname} :: [{name}] <missing>")
                continue
            print(f"\n### {fname} :: [{name}]")
            dump(wb[name])
        wb.close()


if __name__ == "__main__":
    sys.stdout.reconfigure(encoding="utf-8")
    main()

"""Inspect A1-17 template structure."""
import openpyxl

wb = openpyxl.load_workbook(r"wp_templates/A/A1-17 对应数据.xlsx", read_only=True, data_only=True)
print("Sheets:", wb.sheetnames)
for sn in wb.sheetnames:
    ws = wb[sn]
    print(f"\n--- Sheet: {sn} (rows={ws.max_row}, cols={ws.max_column}) ---")
    for i, row in enumerate(ws.iter_rows(min_row=1, max_row=min(35, ws.max_row), values_only=True), 1):
        cells = [str(c)[:50] if c else "" for c in row[:10]]
        line = " | ".join(cells)
        print(f"  R{i:02d}: {line}")
wb.close()

"""Scan docx templates for structure and placeholders (audit-xlsx docx tasks)."""
from __future__ import annotations

import json
import re
from pathlib import Path

from docx import Document
from docx.enum.text import WD_COLOR_INDEX

ROOT = Path(__file__).resolve().parents[2]
TEMPLATES = ROOT / "backend" / "wp_templates" / "A"
OUT = ROOT / "backend" / "data" / "a7_a15_xlsx_audit.json"

BRACKET_RE = re.compile(r"【([^】]+)】")
XX_RE = re.compile(r"(XX|××|×{2,}|201X|202X|20X\d|ABC公司|XX公司|××公司)")
CURLY_RE = re.compile(r"\{\{[^}]+\}\}")


def resolve_template_path(wp_code: str) -> Path:
    """Match by wp_code prefix; tolerates missing space after hyphen (e.g. A9-1向管理层…)."""
    if "-" in wp_code:
        matches = sorted(
            p
            for p in TEMPLATES.iterdir()
            if p.name.startswith(f"{wp_code} ") or p.name.startswith(f"{wp_code}")
            and not p.name.startswith(f"{wp_code}-")
        )
        # exclude A9-1 matching A9-10 etc.
        matches = [p for p in matches if re.match(rf"^{re.escape(wp_code)}(\s|[^-\d])", p.name)]
    else:
        matches = sorted(
            p
            for p in TEMPLATES.iterdir()
            if p.name.startswith(f"{wp_code} ")
            and not re.match(rf"^{re.escape(wp_code)}-\d", p.name)
        )
    if not matches:
        raise FileNotFoundError(f"No template for {wp_code} under {TEMPLATES}")
    if len(matches) > 1:
        xlsx = [p for p in matches if p.suffix.lower() == ".xlsx"]
        return xlsx[0] if xlsx else matches[0]
    return matches[0]


def run_color(run) -> str | None:
    if run.font.highlight_color and run.font.highlight_color != WD_COLOR_INDEX.AUTO:
        return str(run.font.highlight_color)
    if run.font.color and run.font.color.rgb:
        return f"rgb:{run.font.color.rgb}"
    return None


def scan_paragraphs(doc: Document) -> list[dict]:
    items = []
    for i, para in enumerate(doc.paragraphs, 1):
        text = para.text.strip()
        if not text:
            continue
        colors = sorted({c for run in para.runs if (c := run_color(run))})
        items.append(
            {
                "line": i,
                "style": para.style.name if para.style else None,
                "text": text[:300],
                "brackets": BRACKET_RE.findall(text),
                "xx_patterns": XX_RE.findall(text),
                "curly": CURLY_RE.findall(text),
                "run_colors": colors,
            }
        )
    return items


def scan_tables(doc: Document) -> list[dict]:
    tables = []
    for ti, table in enumerate(doc.tables, 1):
        rows_data = []
        for ri, row in enumerate(table.rows):
            cells = []
            for cell in row.cells:
                t = cell.text.strip().replace("\n", " ")[:200]
                if t:
                    cells.append(t)
            if cells:
                rows_data.append({"row": ri + 1, "cells": cells})
        tables.append({"table_index": ti, "row_count": len(table.rows), "col_count": len(table.columns), "rows": rows_data[:30]})
    return tables


def audit_docx(wp_code: str) -> dict:
    path = resolve_template_path(wp_code)
    doc = Document(str(path))
    paragraphs = scan_paragraphs(doc)
    tables = scan_tables(doc)
    all_brackets = []
    all_xx = []
    for p in paragraphs:
        all_brackets.extend(p["brackets"])
        all_xx.extend(p["xx_patterns"])
    for t in tables:
        for row in t["rows"]:
            for cell in row["cells"]:
                all_brackets.extend(BRACKET_RE.findall(cell))
                all_xx.extend(XX_RE.findall(cell))
    return {
        "wp_code": wp_code,
        "filename": path.name,
        "format": "docx",
        "runtime": "wp-popup-docx",
        "paragraph_count": len([p for p in doc.paragraphs if p.text.strip()]),
        "table_count": len(doc.tables),
        "placeholders": {
            "bracket_fields": sorted(set(all_brackets), key=len),
            "xx_date_entity": sorted(set(all_xx)),
        },
        "paragraphs": paragraphs,
        "tables": tables,
        "parser_ready": True,
        "notes": "Popup docx; prefilled-download + OnlyOffice; blue/bracket = guidance per A17/A18",
    }


def merge_audit(entry: dict) -> None:
    existing: list[dict] = []
    if OUT.exists():
        existing = json.loads(OUT.read_text(encoding="utf-8"))
    existing = [e for e in existing if e.get("wp_code") != entry["wp_code"]]
    existing.append(entry)
    existing.sort(key=lambda x: x["wp_code"])
    OUT.write_text(json.dumps(existing, ensure_ascii=False, indent=2), encoding="utf-8")


if __name__ == "__main__":
    import sys

    for code in sys.argv[1:] or ["A8-1", "A8-2"]:
        entry = audit_docx(code)
        merge_audit(entry)
        print(json.dumps({k: entry[k] for k in ("wp_code", "filename", "paragraph_count", "table_count", "placeholders")}, ensure_ascii=False, indent=2))

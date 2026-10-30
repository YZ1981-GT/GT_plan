"""ACNR 目录单元格坐标复核（spec wp-deeplink-locate-coverage-closure 需求 1，design §2）。

13 份旧种子（`backend/data/*_address_registry_seed.json`，本脚本写的 `t_` 种子除外）登记的单元格，扣除既有
cell_verdicts（T05 孤儿 / T06 种子名不一致）已判定的，逐格读模板现算证据并按 design §2.1 处置：
- E1-14 C10 / C15 两格按 `EXPLICIT` 改指模板里的对应格；
- purpose=conclusion 且所在 tab 恰有一个「审计结论」标签格的，改指该标签格；
- 其余停用、replaced_by 父 sheet。
改指 = 新坐标写进 `t_address_registry_seed.json`（带 rule 与模板证据），原格停用、replaced_by 新格；停用写进
`sources/sheet_verdicts.json` 的 cell_verdicts（kind=seed_coordinate_mismatch）。只依赖标准库与 openpyxl。

用法（cwd 任意）：
    python backend/scripts/acnr/review_cell_coordinates.py --write   # 现算并写出（内容不变不写）
    python backend/scripts/acnr/review_cell_coordinates.py --check   # 现算并比对，漂移退出 2 并点名
"""
from __future__ import annotations

import argparse
import json
import re
import sys
import zipfile
from pathlib import Path
from typing import Any

import openpyxl
from openpyxl.utils.cell import column_index_from_string, coordinate_from_string, get_column_letter

_BACKEND = Path(__file__).resolve().parents[2]
_REPO = _BACKEND.parent
_DATA = _BACKEND / "data"
_ACNR = _DATA / "acnr"
_TPL = _BACKEND / "wp_templates"
CATALOG = _ACNR / "global_catalog.json"
VERDICTS = _ACNR / "sources" / "sheet_verdicts.json"
T_SEED = _DATA / "t_address_registry_seed.json"
KIND = "seed_coordinate_mismatch"
SPEC = "wp-deeplink-locate-coverage-closure"
_SEED_RE = re.compile(r"^[a-z]\d*_address_registry_seed\.json$")
#: 「审计结论」标签格：去掉全部空白后整格文字为（可带「四、」类序号与冒号）
CONCLUSION_RE = re.compile(r"^(?:[一二三四五六七八九十]+、)?审计结论[:：]?$")
CONCLUSION_LABEL = "审计结论"
#: 显式改指（design §2.1）：旧 addr_id → 模板里的对应格；row_label / column_header 是模板现值，守卫逐格复算
EXPLICIT: dict[str, dict[str, Any]] = {
    "E1/E1-14/C10": {"a1": "I18", "label": "货币资金合计本期变动率", "purpose": "ratio_analysis", "header_row": 14},
    "E1/E1-14/C15": {"a1": "C34", "label": "银行存款÷资产总额（本期）", "purpose": "ratio_analysis", "header_row": 33},
}
_CLASS_ZH = {"covered": "落在合并区域内", "empty": "空格", "text": "文字格", "value": "取值格",
             "out_of_range": "超出模板已用范围", "no_template": "找不到模板"}
_WORD = re.compile(r"[\u4e00-\u9fffA-Za-z0-9]")
_DOC_NOTE = ("；kind=seed_coordinate_mismatch 为 wp-deeplink-locate-coverage-closure 逐格对模板复核后停用的种子坐标"
             "（backend/scripts/acnr/review_cell_coordinates.py 生成，--check 现算比对）")


def _rel(p: Path) -> str:
    return p.resolve().relative_to(_REPO.resolve()).as_posix()


def _short(v: Any, n: int) -> str:
    s = re.sub(r"\s+", " ", str(v)).strip()
    return s if len(s) <= n else s[: n - 1] + "…"


def _grams(text: str) -> set[str]:
    s = "".join(_WORD.findall(text or ""))
    return {s[i:i + 2] for i in range(len(s) - 1)} if len(s) > 1 else ({s} if s else set())


def cover(label: str, text: str) -> float:
    """语义标签二字片段出现在 text 里的比例（design §1 复核口径）。"""
    g = _grams(label)
    return round(len(g & _grams(text)) / len(g), 2) if g else 0.0


def _is_text(v: Any) -> bool:
    return isinstance(v, str) and v.strip() != "" and not v.startswith("=")


def _parent_code(wp_code: str) -> str:
    m = re.match(r"^([A-S]\d+)", wp_code)
    return m.group(1) if m else wp_code


class Templates:
    """模板册索引（按 workbook.xml 的 tab 名）与懒加载工作簿缓存。"""

    def __init__(self, root: Path = _TPL) -> None:
        self.tabs: dict[Path, list[str]] = {}
        for f in sorted(list(root.rglob("*.xlsx")) + list(root.rglob("*.xlsm"))):
            try:
                with zipfile.ZipFile(f) as z:
                    xml = z.read("xl/workbook.xml").decode("utf-8", "replace")
            except (zipfile.BadZipFile, KeyError, OSError):
                continue
            self.tabs[f] = [n.replace("&amp;", "&") for n in re.findall(r'<sheet[^>]*?name="([^"]+)"', xml)]
        self._books: dict[Path, Any] = {}

    def find(self, parent: str, tab: str) -> Path | None:
        hits = [f for f, names in self.tabs.items() if tab in names]
        pref = [f for f in hits if f.name.startswith(parent)]
        return (pref or hits or [None])[0]

    def sheet(self, book: Path, tab: str):
        if book not in self._books:
            self._books[book] = openpyxl.load_workbook(book, data_only=False)
        return self._books[book][tab]


def _merged_at(ws, r: int, c: int):
    return next((m for m in ws.merged_cells.ranges if m.min_row <= r <= m.max_row and m.min_col <= c <= m.max_col), None)


def _value_at(ws, r: int, c: int) -> Any:
    m = _merged_at(ws, r, c)
    return ws.cell(row=m.min_row, column=m.min_col).value if m else ws.cell(row=r, column=c).value


def cell_evidence(ws, a1: str, label: str) -> dict[str, Any]:
    """一个坐标在模板里的现算证据：类别、值、同行文字、列头、覆盖率（design §1）。"""
    col_letter, r = coordinate_from_string(a1)
    col = column_index_from_string(col_letter)
    max_row, max_col = ws.max_row or 0, ws.max_column or 0
    if r > max_row or col > max_col:
        klass = "out_of_range"
    else:
        m = _merged_at(ws, r, col)
        own = ws.cell(row=r, column=col).value
        if m is not None and (m.min_row, m.min_col) != (r, col):
            klass = "covered"
        elif own is None or str(own).strip() == "":
            klass = "empty"
        elif _is_text(own):
            klass = "text"
        else:
            klass = "value"
    value = _value_at(ws, r, col) if klass != "out_of_range" else None
    row_texts = [f"{get_column_letter(c)}:{_short(v, 30)}" for c in range(1, max_col + 1)
                 if _is_text(v := ws.cell(row=r, column=c).value)] if r <= max_row else []
    header: list[str] = []
    for rr in range(min(r, max_row + 1) - 1, max(0, r - 15), -1):
        v = _value_at(ws, rr, col) if col <= max_col else None
        if _is_text(v):
            header.append(f"{get_column_letter(col)}{rr}:{_short(v, 24)}")
            if len(header) >= 2:
                break
    ctx = " ".join([str(value or "")] + row_texts + header)
    return {"class": klass, "cell_value": None if value is None else _short(value, 60),
            "row_texts": row_texts[:6], "header": header, "ctx_cover": cover(label, ctx)}


def conclusion_labels(ws) -> list[str]:
    """tab 里整格文字为「审计结论」标签的格（design §2.1）。"""
    out = []
    for row in ws.iter_rows():
        for c in row:
            if isinstance(c.value, str) and CONCLUSION_RE.match(re.sub(r"\s+", "", c.value)):
                out.append(c.coordinate)
    return out


def row_label(ws, a1: str) -> str | None:
    """该格所在行第一处文字（行标签）。"""
    _col, r = coordinate_from_string(a1)
    for c in range(1, (ws.max_column or 0) + 1):
        v = _value_at(ws, r, c)
        if isinstance(v, str) and v.strip():
            return _short(v, 40)
    return None


def column_header(ws, a1: str, header_row: int) -> str | None:
    """该格所在列在表头行的值（合并区域取左上角，公式原样）。"""
    col_letter, _r = coordinate_from_string(a1)
    v = _value_at(ws, header_row, column_index_from_string(col_letter))
    return None if v in (None, "") else _short(v, 40)


class ReviewError(RuntimeError):
    """复核输入不成立（父 sheet 不在目录、改指目标与已有 addr_id 撞车等），不写文件。"""


def load_seed_cells() -> list[dict[str, Any]]:
    """13 份旧种子里带 A1 的单元格（与 loaders.from_address_seeds 同口径构造 addr_id）。"""
    out = []
    for p in sorted(_DATA.iterdir()):
        if not (p.is_file() and _SEED_RE.match(p.name)) or p.name == T_SEED.name:
            continue
        for entry in json.loads(p.read_text(encoding="utf-8")).get("entries", []):
            code = entry.get("wp_code", "")
            parent = _parent_code(code)
            for coord in entry.get("coordinates", []):
                a1 = coord.get("cell_address") or coord.get("cell_ref")
                if not a1 or coord.get("semantic_only"):
                    continue
                out.append({"addr_id": f"{parent}/{code}/{a1}", "sheet_addr_id": f"{parent}/{code}", "parent": parent,
                            "sheet_code": code, "a1": a1, "purpose": coord.get("purpose", ""),
                            "label": coord.get("description", "") or coord.get("label", "")})
    return out


def _target_evidence(ws, book: Path, tab: str, a1: str, header_row: int | None) -> dict[str, Any]:
    col_letter, r = coordinate_from_string(a1)
    value = _value_at(ws, r, column_index_from_string(col_letter))
    ev: dict[str, Any] = {"template_file": _rel(book), "template_tab": tab,
                          "cell_value": None if value is None else _short(value, 60), "row_label": row_label(ws, a1)}
    if header_row is not None:
        ev.update({"header_row": header_row, "column_header": column_header(ws, a1, header_row)})
    return ev


def build(templates: Templates | None = None) -> tuple[list[dict[str, Any]], dict[str, Any]]:
    """现算：(本 kind 的停用判定（按 addr_id 排序）, t_ 种子内容)。纯计算，不写文件。"""
    templates = templates or Templates()
    catalog = json.loads(CATALOG.read_text(encoding="utf-8"))
    sheets = {s["addr_id"]: s for s in catalog["sheets"]}
    raw = json.loads(VERDICTS.read_text(encoding="utf-8"))
    prior = {v["addr_id"] for v in raw["cell_verdicts"] if v.get("kind") != KIND}
    seed_cells = load_seed_cells()
    old_ids = {c["addr_id"] for c in seed_cells}
    by_sheet: dict[str, list[dict[str, Any]]] = {}
    for c in seed_cells:
        if c["addr_id"] not in prior:
            by_sheet.setdefault(c["sheet_addr_id"], []).append(c)
    verdicts: list[dict[str, Any]] = []
    added: dict[str, dict[str, dict[str, Any]]] = {}
    for sheet_addr in sorted(by_sheet):
        sheet = sheets.get(sheet_addr)
        if sheet is None or sheet.get("deprecated"):
            raise ReviewError(f"{sheet_addr} 不在目录或已停用（孤儿应已由 T05 判定覆盖）")
        tab = sheet.get("sheet_name") or ""
        book = templates.find(by_sheet[sheet_addr][0]["parent"], tab) if tab else None
        ws = templates.sheet(book, tab) if book else None
        labels = conclusion_labels(ws) if ws is not None else []
        for c in sorted(by_sheet[sheet_addr], key=lambda x: x["addr_id"]):
            ev: dict[str, Any] = {"template_file": _rel(book) if book else None, "template_tab": tab if book else None}
            ev.update(cell_evidence(ws, c["a1"], c["label"]) if ws is not None else {"class": "no_template"})
            klass = _CLASS_ZH.get(ev["class"], ev["class"])
            explicit = EXPLICIT.get(c["addr_id"])
            if ws is not None and explicit:
                target, rule = dict(explicit), "explicit"
            elif ws is not None and c["purpose"] == "conclusion" and len(labels) == 1:
                target, rule = {"a1": labels[0], "label": CONCLUSION_LABEL, "purpose": "conclusion"}, "conclusion_label"
            else:
                target, rule = None, None
            if target is None:
                replaced_by = sheet_addr
                reason = f"种子坐标与模板不符（{klass}）：坐标处内容与「{c['label']}」无关、表内无可改指的格，停用并指向父表"
            else:
                replaced_by = f"{sheet_addr}/{target['a1']}"
                if replaced_by in old_ids:
                    raise ReviewError(f"{c['addr_id']} 的改指目标 {replaced_by} 与旧种子已有坐标撞车，需人工复核")
                tev = _target_evidence(ws, book, tab, target["a1"], target.get("header_row"))
                coord = added.setdefault(sheet_addr, {}).setdefault(target["a1"], {
                    "cell_address": target["a1"], "description": target["label"], "purpose": target["purpose"],
                    "rule": rule, "replaces": [], "evidence": tev})
                coord["replaces"].append(c["addr_id"])
                where = f"{target['a1']}（{tev['row_label']} / {tev.get('column_header')}）" if rule == "explicit" \
                    else f"{target['a1']}「{tev['cell_value']}」"
                reason = f"种子坐标与模板不符（{klass}）：「{c['label']}」在模板对应 {where}，改指新坐标"
            verdicts.append({"addr_id": c["addr_id"], "parent_sheet_code": sheet_addr, "kind": KIND,
                             "verdict": "deprecate", "replaced_by": replaced_by, "reason": reason,
                             "evidence": {**ev, "seed_label": c["label"], "seed_purpose": c["purpose"]}})
    entries = []
    for sheet_addr in sorted(added):
        sheet = sheets[sheet_addr]
        entries.append({"wp_code": sheet["sheet_code"], "sheet_name": sheet["sheet_name"],
                        "coordinates": [added[sheet_addr][a1] for a1 in sorted(added[sheet_addr])]})
    seed = {"description": ("模板核对后的单元格坐标（spec wp-deeplink-locate-coverage-closure，design §2）：每个坐标带改指规则"
                             "（conclusion_label：所在 tab 唯一的「审计结论」标签格；explicit：逐格核对的显式改指）、所替代的旧坐标与"
                             "模板现值证据。由 backend/scripts/acnr/review_cell_coordinates.py --write 生成，--check 现算比对，勿手改。"),
            "source": "backend/scripts/acnr/review_cell_coordinates.py", "entries": entries}
    return verdicts, seed


def _dump(data: Any) -> str:
    return json.dumps(data, ensure_ascii=False, indent=2) + "\n"


def _merged_verdicts(raw: dict[str, Any], verdicts: list[dict[str, Any]]) -> dict[str, Any]:
    out = dict(raw)
    out["cell_verdicts"] = [v for v in raw["cell_verdicts"] if v.get("kind") != KIND] + verdicts
    doc = str(out.get("_doc") or "")
    if KIND not in doc:
        out["_doc"] = doc + _DOC_NOTE
    return out


def _diff(label: str, committed: Any, rebuilt: Any, limit: int = 20) -> list[str]:
    if committed == rebuilt:
        return []
    if isinstance(committed, list) and isinstance(rebuilt, list):
        a = {x.get("addr_id"): x for x in committed if isinstance(x, dict)}
        b = {x.get("addr_id"): x for x in rebuilt if isinstance(x, dict)}
        lines = [f"{label}: 多出 {k}" for k in sorted(set(a) - set(b))] + [f"{label}: 缺少 {k}" for k in sorted(set(b) - set(a))]
        lines += [f"{label}: {k} 内容不同" for k in sorted(set(a) & set(b)) if a[k] != b[k]]
        if not lines:
            lines.append(f"{label}: 顺序不同")
        return lines[:limit]
    return [f"{label}: 内容不同"]


def run(write: bool) -> int:
    verdicts, seed = build()
    raw = json.loads(VERDICTS.read_text(encoding="utf-8"))
    merged = _merged_verdicts(raw, verdicts)
    committed_seed = json.loads(T_SEED.read_text(encoding="utf-8")) if T_SEED.exists() else None
    n_added = sum(len(e["coordinates"]) for e in seed["entries"])
    n_repoint = sum(1 for v in verdicts if v["replaced_by"] != v["parent_sheet_code"])
    summary = f"停用 {len(verdicts)} 格（改指 {n_repoint}、指向父表 {len(verdicts) - n_repoint}），新增坐标 {n_added} 个"
    if write:
        changed = []
        for path, text in ((VERDICTS, _dump(merged)), (T_SEED, _dump(seed))):
            if not path.exists() or path.read_bytes().replace(b"\r\n", b"\n") != text.encode("utf-8"):
                with open(path, "w", encoding="utf-8", newline="\n") as f:
                    f.write(text)
                changed.append(_rel(path))
        print(f"[cell-review] {summary}；" + (f"已写出 {changed}" if changed else "内容未变化，不写文件"))
        return 0
    lines = _diff("cell_verdicts", [v for v in raw["cell_verdicts"] if v.get("kind") == KIND], verdicts)
    lines += _diff("t_address_registry_seed", committed_seed, seed)
    if KIND not in str(raw.get("_doc") or ""):
        lines.append("sheet_verdicts._doc: 缺本 kind 的说明")
    if lines:
        print(f"[cell-review] 漂移 {len(lines)} 处（模板或种子变了未重新生成，或产物被手改）：", file=sys.stderr)
        for line in lines:
            print(f"  {line}", file=sys.stderr)
        return 2
    print(f"[cell-review] OK — {summary}，与模板现算一致")
    return 0


def main(argv: list[str] | None = None) -> int:
    parser = argparse.ArgumentParser(description="ACNR 目录单元格坐标复核（写出 / 比对）")
    mode = parser.add_mutually_exclusive_group(required=True)
    mode.add_argument("--write", action="store_true", help="现算并写出（内容不变不写）")
    mode.add_argument("--check", action="store_true", help="现算并与已提交判定 / 种子比对，漂移退出 2")
    args = parser.parse_args(argv)
    return run(write=args.write)


if __name__ == "__main__":
    sys.exit(main())

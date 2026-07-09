"""Phase0: Extract S3/S15/S20/S21 sheet structure + formula chains for spec validation.
Writes output directly to file to avoid encoding issues."""
from __future__ import annotations
import json
from pathlib import Path
import openpyxl

BASE = Path(
    r"d:\GT_plan\数据\致同通用审计程序及底稿模板（2025年修订）"
    r"\1.致同审计程序及底稿模板（2025年）\6.特定项目程序（S）"
)

OUTPUT = Path(r"d:\GT_plan\.kiro\specs\s-estimate-calculation-workpapers\phase0_raw.json")

TARGETS = {
    "S15": "S15 每股收益和净资产收益率.xlsx",
    "S20": "S20 营业收入扣除情况核查底稿202504.xlsx",
    "S21": "S21 数据资产.xlsx",
    "S3": "S3 会计政策变更、前期差错会计、估计变更2020.xlsx",
}

def extract_formulas(ws):
    formulas = []
    for row in ws.iter_rows():
        for cell in row:
            if isinstance(cell.value, str) and cell.value.startswith("="):
                formulas.append({
                    "cell": f"{cell.column_letter}{cell.row}",
                    "formula": cell.value
                })
    return formulas

def extract_structure(ws, max_rows=40):
    rows_data = []
    for r_i, row in enumerate(ws.iter_rows(values_only=True), start=1):
        if r_i > max_rows:
            break
        cells = []
        for c in row:
            if c is not None and str(c).strip():
                cells.append(str(c).strip().replace("\n", " ")[:150])
        if cells:
            rows_data.append({"row": r_i, "cells": cells})
    return rows_data

def main():
    result = {}
    
    for code, fname in TARGETS.items():
        p = BASE / fname
        if not p.exists():
            print(f"ERROR: {p} not found")
            continue
        wb = openpyxl.load_workbook(p, data_only=False)
        sheets_info = []
        for name in wb.sheetnames:
            ws = wb[name]
            formulas = extract_formulas(ws)
            info = {
                "sheetName": name,
                "rows": ws.max_row or 0,
                "cols": ws.max_column or 0,
                "merged": len(ws.merged_cells.ranges),
                "formulaCount": len(formulas),
                "formulas": formulas[:30],  # limit to first 30
                "structure": extract_structure(ws, max_rows=35),
            }
            sheets_info.append(info)
        result[code] = {
            "file": fname,
            "sheetCount": len(wb.sheetnames),
            "sheetNames": wb.sheetnames,
            "sheets": sheets_info,
        }
        wb.close()
    
    OUTPUT.write_text(json.dumps(result, ensure_ascii=False, indent=2), encoding="utf-8")
    print(f"Done. Output: {OUTPUT}")

if __name__ == "__main__":
    main()

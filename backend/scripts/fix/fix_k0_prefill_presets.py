#!/usr/bin/env python
"""fix_k0_prefill_presets.py — K0 函证公式预设纠偏（幂等）.

spec: k0-confirmation-source-alignment · Task 6 / Requirements 5.1~5.5

改造前 K0 块**整块贴错标签**，三处都不可用：

1. **sheet 名不存在** —— 原写 ``审定表K0-1``，而源模板
   ``backend/wp_templates/K/K0 管理循环函证.xlsx`` 的 10 张可见 sheet 里**没有这个 tab**
   （底稿目录 / 函证程序表K0A / 函证结果汇总表K0-1 / 核实被函证单位信息K0-2 /
   跟函函证过程控制K0-3 / 函证差异调节表K0-4 / 其他应收款替代程序K0-5 /
   其他应付款替代程序K0-6 / 邮件传真回函可靠性验证K0-7 / 函证程序舞弊风险评价表K0-8）。
   函证枢纽压根没有审定表 → 该块永远匹配不上，是死配置。→ 改 ``函证结果汇总表K0-1``。

2. **cell_ref 是审定表口径** —— 原写「期初余额」/「未审数」，而 K0-1 需要取数的位置是
   下区「一、函证情况」矩阵的**「本期（期末）账面金额」两格**（源 ``E29``/``F29``，
   源模板该行**无公式**=手填）。cell_ref 同时是**手工覆盖键** → 写审定表口径的键等于
   让预设永远落不到矩阵上。→ 改为 G0 已确立的矩阵键形态 ``K0-1-matrix-{品种}-book_amount``。

3. **只声明其他应收款** —— 原 ``account_codes=['1221']`` 漏掉其他应付款侧。K0 是
   **两品种**函证（源 ``E28 其他应收款`` / ``F28 其他应付款``）。

取数口径（``report_config`` 只读实证，2026-08-04）::

    其他应收款  BS-009  soe_standalone = TB('1221') - TB('1231-03') + TB('1131')  ← **净额口径**
                        其余三准则     = TB('1221')
    其他应付款  BS-050  *_standalone   = TB('2241') + TB('2231')                  ← 2231 已并入其他应付款
                        *_consolidated = TB('2241')

🔴 **必须按 row_code 精确匹配，不能按 row_name** —— ``BS-075`` 在 soe 侧 row_name 也是
「其他应付款」但 ``formula`` 为 NULL，而在 listed 侧 row_name 竟是「股本」
（同一 row_code 在不同准则下 row_name 不同 = 按名匹配必踩）。

🔴 公式一律 ``PLACEHOLDER``（沿用 H0/G0 已落地的范式），**不写 TB()**：
- BS-009 是**净额**（含备抵扣减），单个 ``TB()`` 表达不出来；
- ``cell_ref`` 是手工覆盖键，写 ``TB()`` 会让「按码取到的 0」伪装成审计师手填值，
  压住语义定位拿到的 ``undefined`` → 「本项目无此科目」与「余额为 0」不可区分；
- 运行态取数一律走 ``four_table`` 语义定位（按科目名在**本项目**科目表定位 + 叶子聚合），
  ``account_codes`` 只作展示与筛选。

用法::

    python backend/scripts/fix/fix_k0_prefill_presets.py --dry-run
    python backend/scripts/fix/fix_k0_prefill_presets.py --check    # exit 1 = 有欠账
    python backend/scripts/fix/fix_k0_prefill_presets.py --apply

🔴 round-trip 自检：先确认 ``json.dumps(json.loads(raw))`` 能**逐字复现**原文，
   不能复现即 exit 2 拒绝写入（防把并发会话的格式/顺序整体重排）。
"""

from __future__ import annotations

import argparse
import json
import sys
from pathlib import Path

REPO_ROOT = Path(__file__).resolve().parents[3]
MAPPING_PATH = REPO_ROOT / "backend" / "data" / "prefill_formula_mapping.json"
K0_XLSX = REPO_ROOT / "backend" / "wp_templates" / "K" / "K0 管理循环函证.xlsx"

WP_CODE = "K0"

#: 源模板真实的函证结果汇总表 tab 名
TARGET_SHEET = "函证结果汇总表K0-1"
OBSOLETE_SHEET = "审定表K0-1"

#: 两品种参考科目码（**仅作展示与筛选**；运行态一律走语义定位，不据此写死取数）。
#: 与 `confirmation/k0-confirmation/k0MatrixSpec.K0_MATRIX_CATEGORIES` 的兜底码并集一致。
K0_REFERENCE_ACCOUNT_CODES = [
    "1221",     # 其他应收款（原值）
    "1231-03",  # 坏账准备-其他应收款（BS-009 soe 口径的减项）
    "1131",     # 应收股利（BS-009 soe 口径的加项）
    "2241",     # 其他应付款
    "2231",     # 应付利息（财会[2018]15 号起并入其他应付款列报）
]

#: 矩阵账面金额两格（源 E29/F29）。cell_ref 形态沿用 G0 已落地的矩阵手工覆盖键，
#: **不用中文 label 构键之外的第二套形态**，也不用裸单元格地址（那不是手工覆盖键）。
TARGET_CELLS = [
    {
        "cell_ref": "K0-1-matrix-其他应收款-book_amount",
        "formula": "=PLACEHOLDER('其他应收款本期（期末）账面金额')",
        "formula_type": "PLACEHOLDER",
        "description": (
            "品种「其他应收款」本期（期末）账面金额（源模板 函证结果汇总表K0-1!E29，该行无公式=手填）。"
            "取数真源 = 相邻 K1 审定表 render-config 的 `project_context.tb_amount`"
            "（走 semantic_account_resolver 按科目名在**本项目**科目表定位 + 叶子聚合，期末余额口径）。"
            "报表行 BS-009 —— soe_standalone 是**净额口径** "
            "TB('1221')-TB('1231-03')+TB('1131')，其余三准则为 TB('1221')；"
            "参考科目码 1221/1231-03/1131 **仅作展示与筛选，运行态不据此取数**（R4.5）。"
            "纠偏：原写 =TB('1221','期初余额') 且 cell_ref 是审定表口径的「期初余额」—— "
            "①函证枢纽无审定表，sheet 名『审定表K0-1』在源 xlsx 不存在；"
            "②单个 TB() 表达不出 BS-009 的净额口径（漏减备抵 1231-03、漏加 1131）；"
            "③cell_ref 是手工覆盖键，写 TB() 会让按码取到的 0 伪装成审计师手填值，"
            "压住语义定位拿到的 undefined → 「本项目无此科目」与「余额为 0」不可区分（R4.4）。"
        ),
    },
    {
        "cell_ref": "K0-1-matrix-其他应付款-book_amount",
        "formula": "=PLACEHOLDER('其他应付款本期（期末）账面金额')",
        "formula_type": "PLACEHOLDER",
        "description": (
            "品种「其他应付款」本期（期末）账面金额（源模板 函证结果汇总表K0-1!F29，该行无公式=手填）。"
            "取数真源 = 相邻 K3 审定表 render-config 的 `project_context.tb_amount`"
            "（走 semantic_account_resolver 按科目名在**本项目**科目表定位 + 叶子聚合，期末余额口径）。"
            "报表行 **BS-050** = TB('2241')+TB('2231')（*_standalone；*_consolidated 仅 2241）；"
            "🔴 必须按 row_code 精确匹配 —— BS-075 在 soe 侧 row_name 也叫「其他应付款」但 formula 为 NULL，"
            "在 listed 侧 row_name 竟是「股本」，按 row_name 匹配必踩（同 K2 的 BS-014/BS-017 同名坑）。"
            "参考科目码 2241/2231 **仅作展示与筛选，运行态不据此取数**（R4.5）。"
        ),
    },
]

#: 只扫这些**语义字段**做校验；`description`/`notes` 会如实写出被纠正的反例
#: （如「原写 =TB('1221','期初余额') 于 审定表K0-1」），把它们纳入比对会让
#: 「说明文字被数成真实引用」，也会让手工补充说明被判成欠账（R5.4）。
SEMANTIC_CELL_FIELDS = ("cell_ref", "formula", "formula_type", "applies_when")


def _cells_semantic(cells) -> list[dict]:
    """cells 的语义视图（**不含 description/notes**）。"""
    return [
        {k: c.get(k) for k in SEMANTIC_CELL_FIELDS if k in c or k in ("cell_ref", "formula", "formula_type")}
        for c in (cells or [])
    ]


def _load_raw() -> str:
    return MAPPING_PATH.read_text(encoding="utf-8")


def _roundtrip_ok(raw: str) -> bool:
    """json.dumps 能否逐字复现原文（否则写盘会整体重排，拒绝执行）。"""
    data = json.loads(raw)
    return json.dumps(data, ensure_ascii=False, indent=2) == raw.rstrip("\n")


def _dump(data: dict) -> str:
    return json.dumps(data, ensure_ascii=False, indent=2) + "\n"


def _plan(data: dict) -> tuple[list[str], dict | None]:
    """返回 (变更清单, K0 块引用)。清单为空 = 已达目标状态（幂等）。"""
    blocks = [m for m in data["mappings"] if str(m.get("wp_code", "")) == WP_CODE]
    if not blocks:
        return ["[ERROR] prefill_formula_mapping.json 无 K0 块"], None
    if len(blocks) > 1:
        return [f"[ERROR] K0 块重复 {len(blocks)} 个，需人工核对"], None
    blk = blocks[0]

    changes: list[str] = []
    if blk.get("sheet") != TARGET_SHEET:
        changes.append(f"sheet: 「{blk.get('sheet')}」→「{TARGET_SHEET}」")
    if blk.get("account_codes") != K0_REFERENCE_ACCOUNT_CODES:
        changes.append(
            f"account_codes: {blk.get('account_codes')} → 两品种参考码"
            f"（{len(K0_REFERENCE_ACCOUNT_CODES)} 项，仅作展示与筛选）"
        )
    if _cells_semantic(blk.get("cells")) != _cells_semantic(TARGET_CELLS):
        old = [f"{c.get('cell_ref')}={c.get('formula')}" for c in blk.get("cells") or []]
        changes.append(f"cells: {old} → 2 条矩阵账面金额 PLACEHOLDER（取数改由语义定位下发）")
    # description 不参与语义比对，但不许为空（溯源可读性红线）
    empty_desc = [c.get("cell_ref") for c in blk.get("cells") or [] if not str(c.get("description") or "").strip()]
    if empty_desc:
        changes.append(f"description 为空: {empty_desc}")
    return changes, blk


def _semantic_view(blk: dict) -> str:
    """只取语义字段的规范化视图（R5.4：校验器不扫 description/notes）。"""
    view = {k: blk.get(k) for k in ("wp_code", "sheet", "account_codes")}
    view["cells"] = _cells_semantic(blk.get("cells"))
    return json.dumps(view, ensure_ascii=False, sort_keys=True)


def _verify_target_sheet_exists() -> str | None:
    """校验目标 sheet 名确实在源模板中（防把一个错名换成另一个错名）。"""
    try:
        import openpyxl
    except ImportError:
        return None  # 无 openpyxl 时跳过（CI 有）
    if not K0_XLSX.exists():
        return f"源模板不存在: {K0_XLSX}"
    wb = openpyxl.load_workbook(K0_XLSX)
    visible = {n for n in wb.sheetnames if wb[n].sheet_state == "visible"}
    if TARGET_SHEET not in visible:
        return f"目标 sheet「{TARGET_SHEET}」不在源模板可见 sheet 中：{sorted(visible)}"
    if OBSOLETE_SHEET in set(wb.sheetnames):
        return f"源模板竟存在「{OBSOLETE_SHEET}」→ 本脚本前提失效，请重新核对"
    return None


def main() -> int:
    ap = argparse.ArgumentParser(description="K0 函证公式预设纠偏（幂等）")
    g = ap.add_mutually_exclusive_group()
    g.add_argument("--dry-run", action="store_true", help="仅打印变更计划（默认）")
    g.add_argument("--check", action="store_true", help="有欠账则 exit 1")
    g.add_argument("--apply", action="store_true", help="写入")
    args = ap.parse_args()

    raw = _load_raw()
    if not _roundtrip_ok(raw):
        print(
            "[FATAL] prefill_formula_mapping.json 无法 round-trip 复现"
            "（缩进/键序与 json.dumps(indent=2) 不一致）→ 拒绝写入以免整体重排。",
            file=sys.stderr,
        )
        return 2

    err = _verify_target_sheet_exists()
    if err:
        print(f"[FATAL] {err}", file=sys.stderr)
        return 2

    data = json.loads(raw)
    total_before = len(data["mappings"])
    changes, blk = _plan(data)

    if blk is None:
        for c in changes:
            print(c, file=sys.stderr)
        return 2

    if not changes:
        print("[OK] K0 预设已达目标状态，0 项欠账")
        return 0

    print(f"[PLAN] K0 预设待修 {len(changes)} 项：")
    for c in changes:
        print(f"  - {c}")

    if args.check:
        return 1
    if not args.apply:
        print("（--dry-run，未写入；加 --apply 执行）")
        return 0

    blk["sheet"] = TARGET_SHEET
    blk["account_codes"] = list(K0_REFERENCE_ACCOUNT_CODES)
    blk["cells"] = json.loads(json.dumps(TARGET_CELLS, ensure_ascii=False))

    # 只触碰 K0 块：块数不变，其余块逐字节不变
    assert len(data["mappings"]) == total_before, "块数变了"
    old_data = json.loads(raw)
    for i, (old_blk, new_blk) in enumerate(zip(old_data["mappings"], data["mappings"])):
        if str(new_blk.get("wp_code")) == WP_CODE:
            continue
        assert json.dumps(old_blk, ensure_ascii=False, sort_keys=True) == json.dumps(
            new_blk, ensure_ascii=False, sort_keys=True
        ), f"非 K0 块 #{i}({new_blk.get('wp_code')}) 被改动"

    MAPPING_PATH.write_text(_dump(data), encoding="utf-8")
    print(f"[APPLIED] 已写入 {MAPPING_PATH}（{len(changes)} 项）")

    # 写后自检：重新读一遍应为 0 欠账（幂等）
    changes2, blk2 = _plan(json.loads(_load_raw()))
    if changes2:
        print(f"[FATAL] 写后仍有欠账: {changes2}", file=sys.stderr)
        return 2
    assert blk2 is not None and _semantic_view(blk2) == _semantic_view(blk)
    print("[OK] 写后复检 0 项欠账（幂等）")
    return 0


if __name__ == "__main__":
    sys.exit(main())

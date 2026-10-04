#!/usr/bin/env python
"""补齐 F3/F4/F5 公式预设 + 纠错 F2/F5 已有预设（幂等）。

**改造前现状**（2026-08-03 实测）：
- F3：审定表 / 明细表两个骨架块 items=0，完全无预设公式
- F4：审定表 / 明细表两个骨架块 items=0
- F5：审定表一个骨架块 items=0
- F2：已有 18 个骨架块 items=0（`fix_f2_prefill_presets.py` 已处理 F2-1/F2-2/F2-18/F2-47/
  附注两块共 85 条），本脚本只处理 **F2 的科目语义纠错**（5.3/5.4/5.5）

**本脚本新增/修正**：
1. F3 审定表块：TB('2201') 期初/期末 + WP('F3','明细表F3-2',...) 四条联动
2. F3 明细表块：TB('2201') 两条 + AUX 占位一条
3. F3 披露两块：各含 WP + TB 合计核对
4. F4 审定表块：TB('2202') 期初/期末 + WP 联动（按性质 + 按账龄 + 长期挂账）
5. F4 明细表块：TB('2202') + AUX 占位（禁反引审定表）
6. F4 披露两块：WP('F4','长期挂账检查表F4-5',...) + TB 合计
7. F5 审定表块：TB('6401~6499','本期发生额') + WP 联动三条
8. F2 纠错：1402 描述纠正 / 1403 描述纠正 / 审定表口径统一 1401~1499
9. F3/F4 虚构 AUX 维度值删除（'TOP1' 等），改 XX单位 占位

**硬约束**

- `page_key = f"workpaper:{wp_code}"` 忽略 sheet → cell_ref 全局唯一
- 明细表块禁 WP()（防 审定表↔明细表 成环）
- 损益类无「期初余额/期末余额」概念，只有「本期发生额」

科目口径依据（DB 只读实证）：

- `BS-044 应付票据 = TB('2201','期末余额')` 四准则一致
- `BS-045 应付账款 = TB('2202','期末余额')` 四准则一致
- `IS-002 营业成本 = SUM_TB('6401~6499','本期发生额')` 四准则一致

Usage::

    python backend/scripts/fix/fix_f_cycle_prefill_presets.py --dry-run
    python backend/scripts/fix/fix_f_cycle_prefill_presets.py
    python backend/scripts/fix/fix_f_cycle_prefill_presets.py --check

spec: .kiro/specs/f-cycle-four-table-extraction-and-disclosure-completion/
      Requirements 5.1~5.8 / Property 8, 9
"""
from __future__ import annotations

import argparse
import json
import sys
from pathlib import Path
from typing import Any

_BACKEND = Path(__file__).resolve().parent.parent.parent
MAPPING_PATH = _BACKEND / "data" / "prefill_formula_mapping.json"

# ═══════════════════════════════════ F3 应付票据 ═══════════════════════════════

F3_SHEET_ADJ = "审定表F3-1"
F3_SHEET_DETAIL = "明细表F3-2"
F3_SHEET_DISC_LISTED = "附注披露信息(上市公司)"
F3_SHEET_DISC_SOE = "附注披露信息(国企)"

F3_ADJ_CELLS: list[dict[str, Any]] = [
    {
        "cell_ref": "期初余额",
        "formula": "=TB('2201','期初余额')",
        "formula_type": "TB",
        "description": "应付票据期初余额（报表行 BS-044 口径）",
    },
    {
        "cell_ref": "期末未审数",
        "formula": "=TB('2201','期末余额')",
        "formula_type": "TB",
        "description": "应付票据期末未审余额",
    },
    {
        "cell_ref": "上年审定数",
        "formula": "=PREV('F3','审定表F3-1','期末审定数')",
        "formula_type": "PREV",
        "description": "上年 F3-1 审定表期末审定数（期初核对基准）",
    },
    {
        "cell_ref": "明细合计期末",
        "formula": "=WP('F3','明细表F3-2','期末审定余额合计')",
        "formula_type": "WP",
        "description": "审定表期末数来源 = F3-2 明细表期末审定余额合计（源模板 F3-1 B/F/G/H/I 列 SUMPRODUCT 自 F3-2）",
    },
    {
        "cell_ref": "明细合计期初",
        "formula": "=WP('F3','明细表F3-2','期初审定余额合计')",
        "formula_type": "WP",
        "description": "审定表期初数来源 = F3-2 明细表期初审定余额合计",
    },
    {
        "cell_ref": "明细银行承兑合计",
        "formula": "=WP('F3','明细表F3-2','银行承兑汇票审定余额合计')",
        "formula_type": "WP",
        "description": "审定表银行承兑行 = F3-2 按票据种类筛选的银行承兑汇票合计",
    },
    {
        "cell_ref": "明细商业承兑合计",
        "formula": "=WP('F3','明细表F3-2','商业承兑汇票审定余额合计')",
        "formula_type": "WP",
        "description": "审定表商业承兑行 = F3-2 按票据种类筛选的商业承兑汇票合计",
    },
]

F3_DETAIL_BLOCK: dict[str, Any] = {
    "wp_code": "F3",
    "wp_name": "应付票据明细表",
    "sheet": F3_SHEET_DETAIL,
    "account_codes": ["2201"],
    "cells": [
        {
            "cell_ref": "明细期初余额",
            "formula": "=TB('2201','期初余额')",
            "formula_type": "TB",
            "description": "明细表期初未审余额合计核对数（与总账勾稽）",
        },
        {
            "cell_ref": "明细期末余额",
            "formula": "=TB('2201','期末余额')",
            "formula_type": "TB",
            "description": "明细表期末未审余额合计核对数",
        },
        {
            "cell_ref": "明细供应商期末余额",
            "formula": "=AUX('2201','供应商','XX单位','期末余额')",
            "formula_type": "AUX",
            "description": "按供应商（辅助维度）取某单位期末余额（逐行取数走批量导入，本条供单户复核）",
        },
    ],
}

F3_DISC_LISTED_BLOCK: dict[str, Any] = {
    "wp_code": "F3",
    "wp_name": "应付票据附注披露信息（上市公司）",
    "sheet": F3_SHEET_DISC_LISTED,
    "account_codes": ["2201"],
    "cells": [
        {
            "cell_ref": "披露期末合计",
            "formula": "=TB('2201','期末余额')",
            "formula_type": "TB",
            "description": "附注披露表期末合计核对数（= 报表行 BS-044）",
        },
        {
            "cell_ref": "披露审定期末",
            "formula": "=WP('F3','审定表F3-1','期末审定数')",
            "formula_type": "WP",
            "description": "附注披露表期末数来源 = 审定表期末审定数",
        },
    ],
}

F3_DISC_SOE_BLOCK: dict[str, Any] = {
    "wp_code": "F3",
    "wp_name": "应付票据附注披露信息（国企）",
    "sheet": F3_SHEET_DISC_SOE,
    "account_codes": ["2201"],
    "cells": [
        {
            "cell_ref": "披露soe期末合计",
            "formula": "=TB('2201','期末余额')",
            "formula_type": "TB",
            "description": "国企附注披露表期末合计核对数",
        },
        {
            "cell_ref": "披露soe审定期末",
            "formula": "=WP('F3','审定表F3-1','期末审定数')",
            "formula_type": "WP",
            "description": "国企附注披露表期末数来源 = 审定表期末审定数",
        },
    ],
}

# ═══════════════════════════════════ F4 应付账款 ═══════════════════════════════

F4_SHEET_ADJ = "审定表F4-1"
F4_SHEET_DETAIL = "明细表F4-2"
F4_SHEET_DISC_LISTED = "附注披露信息(上市公司)"
F4_SHEET_DISC_SOE = "附注披露信息(国企)"

F4_ADJ_CELLS: list[dict[str, Any]] = [
    {
        "cell_ref": "期初余额",
        "formula": "=TB('2202','期初余额')",
        "formula_type": "TB",
        "description": "应付账款期初余额（报表行 BS-045 口径）",
    },
    {
        "cell_ref": "期末未审数",
        "formula": "=TB('2202','期末余额')",
        "formula_type": "TB",
        "description": "应付账款期末未审余额",
    },
    {
        "cell_ref": "上年审定数",
        "formula": "=PREV('F4','审定表F4-1','期末审定数')",
        "formula_type": "PREV",
        "description": "上年 F4-1 审定表期末审定数",
    },
    {
        "cell_ref": "明细合计期末",
        "formula": "=WP('F4','明细表F4-2','期末审定余额合计')",
        "formula_type": "WP",
        "description": "按性质区期末数 = F4-2 明细表期末审定余额合计（源模板 F/G/H 列 SUMIF 自 F4-2）",
    },
    {
        "cell_ref": "明细合计期初",
        "formula": "=WP('F4','明细表F4-2','期初审定余额合计')",
        "formula_type": "WP",
        "description": "按性质区期初数 = F4-2 明细表期初审定余额合计",
    },
    {
        "cell_ref": "明细按账龄合计",
        "formula": "=WP('F4','明细表F4-2','按账龄汇总合计')",
        "formula_type": "WP",
        "description": "按账龄区期末数 = F4-2 明细表按账龄汇总（源模板引 明细表F4-2!N32:X32）",
    },
    {
        "cell_ref": "长期挂账审定合计",
        "formula": "=WP('F4','长期挂账检查表F4-5','审定余额合计')",
        "formula_type": "WP",
        "description": "账龄超1年重要应付账款审定合计（← F4-5 长期挂账检查表）",
    },
]

F4_DETAIL_BLOCK: dict[str, Any] = {
    "wp_code": "F4",
    "wp_name": "应付账款明细表",
    "sheet": F4_SHEET_DETAIL,
    "account_codes": ["2202"],
    "cells": [
        {
            "cell_ref": "明细期初余额",
            "formula": "=TB('2202','期初余额')",
            "formula_type": "TB",
            "description": "明细表期初未审余额合计核对数（与总账勾稽）",
        },
        {
            "cell_ref": "明细期末余额",
            "formula": "=TB('2202','期末余额')",
            "formula_type": "TB",
            "description": "明细表期末未审余额合计核对数",
        },
        {
            "cell_ref": "明细供应商期末余额",
            "formula": "=AUX('2202','供应商','XX单位','期末余额')",
            "formula_type": "AUX",
            "description": "按供应商（辅助维度）取某单位期末余额（逐行取数走批量导入，本条供单户复核）",
        },
    ],
}

F4_DISC_LISTED_BLOCK: dict[str, Any] = {
    "wp_code": "F4",
    "wp_name": "应付账款附注披露信息（上市公司）",
    "sheet": F4_SHEET_DISC_LISTED,
    "account_codes": ["2202"],
    "cells": [
        {
            "cell_ref": "披露按性质合计",
            "formula": "=TB('2202','期末余额')",
            "formula_type": "TB",
            "description": "上市披露按性质主表合计核对数（= 报表行 BS-045）",
        },
        {
            "cell_ref": "披露长期挂账合计",
            "formula": "=WP('F4','长期挂账检查表F4-5','审定余额合计')",
            "formula_type": "WP",
            "description": "上市披露「账龄超过1年的重要应付账款」合计来源",
        },
    ],
}

F4_DISC_SOE_BLOCK: dict[str, Any] = {
    "wp_code": "F4",
    "wp_name": "应付账款附注披露信息（国企）",
    "sheet": F4_SHEET_DISC_SOE,
    "account_codes": ["2202"],
    "cells": [
        {
            "cell_ref": "披露soe期末合计",
            "formula": "=TB('2202','期末余额')",
            "formula_type": "TB",
            "description": "国企披露按账龄主表合计核对数（= 报表行 BS-045）",
        },
        {
            "cell_ref": "披露soe长期挂账合计",
            "formula": "=WP('F4','长期挂账检查表F4-5','审定余额合计')",
            "formula_type": "WP",
            "description": "国企披露「账龄超过1年的重要应付账款」合计来源",
        },
    ],
}

# ═══════════════════════════════════ F5 营业成本 ═══════════════════════════════

F5_SHEET_ADJ = "营业务成本审定表F5-1"

F5_ADJ_CELLS: list[dict[str, Any]] = [
    {
        "cell_ref": "本期发生额",
        "formula": "=TB_SUM('6401~6499','本期发生额')",
        "formula_type": "TB_SUM",
        "description": (
            "营业成本本期发生额合计（报表行 IS-002 口径；覆盖 6401 主营 + 6402/6404 其他业务成本）。"
            "🔴 损益类无「期初/期末余额」概念，用「本期发生额」。"
        ),
    },
    {
        "cell_ref": "上年发生额",
        "formula": "=PREV('F5','营业务成本审定表F5-1','本期审定发生额')",
        "formula_type": "PREV",
        "description": "上年 F5-1 审定表本期审定发生额（比较数）",
    },
    {
        "cell_ref": "主营明细合计",
        "formula": "=WP('F5','主营业务成本月度明细表F5-2','审定发生额合计')",
        "formula_type": "WP",
        "description": "主营业务成本明细合计（← F5-2，源模板 F5-1 引 F5-2 月度汇总行）",
    },
    {
        "cell_ref": "其他业务成本合计",
        "formula": "=WP('F5','其他业务成本明细表F5-3','审定发生额合计')",
        "formula_type": "WP",
        "description": "其他业务成本明细合计（← F5-3）",
    },
    {
        "cell_ref": "调整分录合计",
        "formula": "=WP('F5','调整分录汇总F5-4','调整分录净额')",
        "formula_type": "WP",
        "description": "AJE/RJE 净调整额（← F5-4 调整分录汇总）",
    },
]

# ═══════════════════════════════ F2 纠错条目 ═══════════════════════════════════

#: F2 科目描述纠错映射 —— 改描述不改公式（公式本身指向的码是对的，只是中文描述写错了）
F2_DESCRIPTION_FIXES: list[dict[str, str]] = [
    # 5.3(a): 1402 是在途物资不是生产成本
    {
        "wp_code": "F2",
        "sheets": "生产成本明细表F2-41,直接人工分析表F2-42,制造费用明细表F2-43",
        "match_formula_contains": "TB('1402'",
        "old_description_contains": "在产品（生产成本）",
        "new_description": "在途物资（1402；源模板引此为「生产成本」科目，标准科目表两变体均为「在途物资」）",
    },
    # 5.4(b): 1403 是原材料不是库存商品
    {
        "wp_code": "F2",
        "sheets": "盘点计划问卷F2-21,监盘计划F2-22,监盘小结F2-23,抽盘结果汇总表F2-25,盘点倒轧表F2-26",
        "match_formula_contains": "'1403'",
        "old_description_contains": "库存商品",
        "new_description": "原材料（1403；标准科目表两变体均为「原材料」，非「库存商品」）",
    },
]

#: F2 审定表口径统一 —— 5.5(c): TB_SUM('1401~1461') → '1401~1499'（BS-010 口径）
F2_RANGE_FIX = {
    "wp_code": "F2",
    "sheet": "审定表F2-1",
    "old_range": "1401~1461",
    "new_range": "1401~1499",
}

#: F5 损益口径纠错 —— 5.1: 旧块 `审定表F5-1` 里的期初/期末余额改本期发生额
F5_PL_PERIOD_FIXES = [
    ("期初余额", "本期发生额"),
    ("期末余额", "本期发生额"),
]

#: F5 旧块 sheet 名纠正 —— 源 xlsx 真实 tab 名
F5_OLD_SHEET = "审定表F5-1"
F5_CORRECT_SHEET = "营业务成本审定表F5-1"

#: F3/F4 虚构 AUX 维度值删除 —— 5.6
F34_AUX_FABRICATIONS = ("TOP1", "TOP2", "TOP3", "TOP4", "TOP5")
#: F2 虚构 AUX 维度值（A类/B类/长库龄/呆滞）—— 也需清理
F2_AUX_FABRICATIONS = ("A类", "B类", "长库龄", "呆滞")


# ═══════════════════════════════ 脚本引擎 ═══════════════════════════════════════


def _blocks(doc: dict[str, Any]) -> list[dict[str, Any]]:
    return doc["mappings"]


def _find_block(doc: dict[str, Any], wp_code: str, sheet: str) -> dict[str, Any] | None:
    for b in _blocks(doc):
        if b.get("wp_code") == wp_code and b.get("sheet") == sheet:
            return b
    return None


def _cells(block: dict[str, Any]) -> list[dict[str, Any]]:
    """块内公式格列表的**唯一**读写口径 —— 恒为 `cells` 键（FC-11 根因修复）。

    🔴 spec: f3-sync-coverage-and-first-canary · Task 17 · Requirement 7.2③（FC-11）

    改造前本脚本有 13 处 `items_key = "items" if "items" in block else "cells"` 的兼容读，
    加上 `_ensure_block` / `_ensure_cells` 两处**主动创建 `items` 键**，共同造成
    「脚本 `--check` 报 exit 0 全部到位，而运行时四个消费方一条都读不到」的矛盾：

    运行时消费方（按值实测，**全部只读 `cells`**）：
      * `wp_template_init_service.py:670`      预填写入 xlsx
      * `formula_management/preset_library.py:163`  公式管理页预设
      * `formula_reverse_index.py:294`         反向索引建边
      * `linkage_graph_builder.py:214`         依赖图

    ⇒ 本 helper 是单一口径：不存在 `cells` 时建空列表并挂上，**永不**创建 `items`。
      `items` 键的存在本身由 `validate()` 判红（防回退）。
    """
    existing = block.get("cells")
    if not isinstance(existing, list):
        existing = []
        block["cells"] = existing
    return existing


def _cell_refs_in_block(block: dict[str, Any]) -> set[str]:
    return {str(c.get("cell_ref") or "") for c in _cells(block)}


def _ensure_block(doc: dict[str, Any], block_def: dict[str, Any]) -> tuple[dict[str, Any], bool]:
    """确保块存在；返回 (block, is_new)。"""
    existing = _find_block(doc, block_def["wp_code"], block_def["sheet"])
    if existing is not None:
        return existing, False
    # 创建新块
    # 🔴 FC-11 根因之二：此处原为 `"items": []` —— 新建块直接就造运行时读不到的键。
    #    spec 的 Requirement 7.2③ 只点了 `_ensure_cells`，漏了这里；两处必须一起修，
    #    否则下次新增块仍然产出死配置。
    new_block = {
        "wp_code": block_def["wp_code"],
        "wp_name": block_def.get("wp_name", ""),
        "sheet": block_def["sheet"],
        "account_codes": block_def.get("account_codes", []),
        "cells": [],
    }
    _blocks(doc).append(new_block)
    return new_block, True


def _ensure_cells(block: dict[str, Any], cells: list[dict[str, Any]]) -> list[str]:
    """确保块中存在所有 cells；返回新增的 cell_ref 列表。

    🔴 FC-11 根因之一（改造前 L417-432）：原实现 `items_key = "items" if "items" in block
       else "cells"`，且块无键时主动 `block["items"] = []` ⇒ 造出运行时读不到的键。
       现只走 `_cells()` 单一口径。
    """
    target = _cells(block)
    existing_refs = {str(c.get("cell_ref") or "") for c in target}
    added: list[str] = []
    for cell in cells:
        ref = cell["cell_ref"]
        if ref not in existing_refs:
            target.append(cell)
            existing_refs.add(ref)
            added.append(ref)
    return added


def _fix_aux_fabrications(doc: dict[str, Any]) -> list[str]:
    """删除 F 类的虚构 AUX 维度值（TOP1~TOP5/A类/B类/长库龄/呆滞），改为 XX单位/XX分类 占位。"""
    changes: list[str] = []
    for b in _blocks(doc):
        wp_code = b.get("wp_code", "")
        if not wp_code.startswith("F"):
            continue
        for cell in _cells(b):
            formula = cell.get("formula", "")
            for fab in F34_AUX_FABRICATIONS:
                if f"'{fab}'" in formula:
                    formula = formula.replace(f"'{fab}'", "'XX单位'")
                    cell["formula"] = formula
                    changes.append(
                        f"{wp_code}/{b.get('sheet','')}: {cell.get('cell_ref','')}: "
                        f"AUX 维度 '{fab}' -> 'XX单位'"
                    )
            for fab in F2_AUX_FABRICATIONS:
                if f"'{fab}'" in formula:
                    formula = formula.replace(f"'{fab}'", "'XX分类'")
                    cell["formula"] = formula
                    changes.append(
                        f"{wp_code}/{b.get('sheet','')}: {cell.get('cell_ref','')}: "
                        f"AUX 维度 '{fab}' -> 'XX分类'"
                    )
    return changes


def _fix_f2_descriptions(doc: dict[str, Any]) -> list[str]:
    """修正 F2 科目描述错误（5.3/5.4）。"""
    changes: list[str] = []
    for fix in F2_DESCRIPTION_FIXES:
        target_sheets = [s.strip() for s in fix["sheets"].split(",")]
        for b in _blocks(doc):
            if b.get("wp_code") != fix["wp_code"]:
                continue
            if b.get("sheet") not in target_sheets:
                continue
            for cell in _cells(b):
                formula = cell.get("formula", "")
                desc = cell.get("description", "")
                if (
                    fix["match_formula_contains"] in formula
                    and fix["old_description_contains"] in desc
                ):
                    cell["description"] = fix["new_description"]
                    changes.append(
                        f"F2/{b.get('sheet','')}: {cell.get('cell_ref','')}: "
                        f"描述纠正 → '{fix['new_description'][:40]}...'"
                    )
    return changes


def _fix_f2_range(doc: dict[str, Any]) -> list[str]:
    """修正 F2 审定表口径 1401~1461 → 1401~1499（5.5）。"""
    changes: list[str] = []
    fix = F2_RANGE_FIX
    block = _find_block(doc, fix["wp_code"], fix["sheet"])
    if block is None:
        return changes
    for cell in _cells(block):
        formula = cell.get("formula", "")
        if fix["old_range"] in formula:
            cell["formula"] = formula.replace(fix["old_range"], fix["new_range"])
            changes.append(
                f"F2/{fix['sheet']}: {cell.get('cell_ref','')}: "
                f"口径 {fix['old_range']} → {fix['new_range']}"
            )
    return changes


def _fix_f5_pl_period(doc: dict[str, Any]) -> list[str]:
    """修正 F5 旧块损益口径：期初/期末余额 → 本期发生额（5.1）。"""
    changes: list[str] = []
    for b in _blocks(doc):
        if b.get("wp_code") != "F5":
            continue
        for cell in _cells(b):
            formula = cell.get("formula", "")
            for old_period, new_period in F5_PL_PERIOD_FIXES:
                if f"'{old_period}'" in formula:
                    cell["formula"] = formula.replace(f"'{old_period}'", f"'{new_period}'")
                    formula = cell["formula"]
                    changes.append(
                        f"F5/{b.get('sheet','')}: {cell.get('cell_ref','')}: "
                        f"损益口径 '{old_period}' → '{new_period}'"
                    )
    return changes


def _merge_f5_old_block(doc: dict[str, Any]) -> list[str]:
    """把 F5 旧块 `审定表F5-1` 的独有 items 合并到正确名称块，然后删旧块。"""
    changes: list[str] = []
    old_block = _find_block(doc, "F5", F5_OLD_SHEET)
    if old_block is None:
        return changes  # 已清理

    new_block = _find_block(doc, "F5", F5_CORRECT_SHEET)
    if new_block is None:
        # 只改 sheet 名
        old_block["sheet"] = F5_CORRECT_SHEET
        changes.append(f"F5: 旧块 sheet 名 '{F5_OLD_SHEET}' → '{F5_CORRECT_SHEET}'")
        return changes

    # 两块都在：把旧块独有公式格合并到新块，然后删旧块
    # 🔴 FC-11 的读写边界：**读**存量旧块要兼容 `items`（那正是待迁移的死配置，
    #    不兼容读就迁不动），**写**新块只走 `_cells()` 单一口径。这与 F5 spec
    #    需求 2.3「契约只声明主键 / 前端 legacy 读回退保留」是同一条原则。
    target = _cells(new_block)
    existing_refs = {str(c.get("cell_ref") or "") for c in target}
    legacy_source = old_block.get("cells") or old_block.get("items") or []

    for cell in legacy_source:
        ref = cell.get("cell_ref", "")
        if ref and ref not in existing_refs:
            target.append(cell)
            existing_refs.add(ref)
            changes.append(f"F5/{F5_CORRECT_SHEET}: 从旧块迁入 {ref}")

    # 删除旧块
    _blocks(doc).remove(old_block)
    changes.append(f"F5: 删除旧块 '{F5_OLD_SHEET}'（已合并到 '{F5_CORRECT_SHEET}'）")
    return changes


def apply(doc: dict[str, Any]) -> list[str]:
    """应用所有改动，返回变更描述列表。"""
    changes: list[str] = []

    # ─── F3 ───────────────────────────────────────────────────────────────────
    adj_block, adj_new = _ensure_block(doc, {"wp_code": "F3", "sheet": F3_SHEET_ADJ,
                                              "wp_name": "应付票据审定表", "account_codes": ["2201"]})
    if adj_new:
        changes.append(f"F3: 新建块 {F3_SHEET_ADJ}")
    added = _ensure_cells(adj_block, F3_ADJ_CELLS)
    for ref in added:
        changes.append(f"F3/{F3_SHEET_ADJ}: 新增 {ref}")

    det_block, det_new = _ensure_block(doc, F3_DETAIL_BLOCK)
    if det_new:
        changes.append(f"F3: 新建块 {F3_SHEET_DETAIL}")
    added = _ensure_cells(det_block, F3_DETAIL_BLOCK["cells"])
    for ref in added:
        changes.append(f"F3/{F3_SHEET_DETAIL}: 新增 {ref}")

    disc_l, disc_l_new = _ensure_block(doc, F3_DISC_LISTED_BLOCK)
    if disc_l_new:
        changes.append(f"F3: 新建块 {F3_SHEET_DISC_LISTED}")
    added = _ensure_cells(disc_l, F3_DISC_LISTED_BLOCK["cells"])
    for ref in added:
        changes.append(f"F3/{F3_SHEET_DISC_LISTED}: 新增 {ref}")

    disc_s, disc_s_new = _ensure_block(doc, F3_DISC_SOE_BLOCK)
    if disc_s_new:
        changes.append(f"F3: 新建块 {F3_SHEET_DISC_SOE}")
    added = _ensure_cells(disc_s, F3_DISC_SOE_BLOCK["cells"])
    for ref in added:
        changes.append(f"F3/{F3_SHEET_DISC_SOE}: 新增 {ref}")

    # ─── F4 ───────────────────────────────────────────────────────────────────
    adj4, adj4_new = _ensure_block(doc, {"wp_code": "F4", "sheet": F4_SHEET_ADJ,
                                          "wp_name": "应付账款审定表", "account_codes": ["2202"]})
    if adj4_new:
        changes.append(f"F4: 新建块 {F4_SHEET_ADJ}")
    added = _ensure_cells(adj4, F4_ADJ_CELLS)
    for ref in added:
        changes.append(f"F4/{F4_SHEET_ADJ}: 新增 {ref}")

    det4, det4_new = _ensure_block(doc, F4_DETAIL_BLOCK)
    if det4_new:
        changes.append(f"F4: 新建块 {F4_SHEET_DETAIL}")
    added = _ensure_cells(det4, F4_DETAIL_BLOCK["cells"])
    for ref in added:
        changes.append(f"F4/{F4_SHEET_DETAIL}: 新增 {ref}")

    disc4_l, disc4_l_new = _ensure_block(doc, F4_DISC_LISTED_BLOCK)
    if disc4_l_new:
        changes.append(f"F4: 新建块 {F4_SHEET_DISC_LISTED}")
    added = _ensure_cells(disc4_l, F4_DISC_LISTED_BLOCK["cells"])
    for ref in added:
        changes.append(f"F4/{F4_SHEET_DISC_LISTED}: 新增 {ref}")

    disc4_s, disc4_s_new = _ensure_block(doc, F4_DISC_SOE_BLOCK)
    if disc4_s_new:
        changes.append(f"F4: 新建块 {F4_SHEET_DISC_SOE}")
    added = _ensure_cells(disc4_s, F4_DISC_SOE_BLOCK["cells"])
    for ref in added:
        changes.append(f"F4/{F4_SHEET_DISC_SOE}: 新增 {ref}")

    # ─── F5 ───────────────────────────────────────────────────────────────────
    adj5, adj5_new = _ensure_block(doc, {"wp_code": "F5", "sheet": F5_SHEET_ADJ,
                                          "wp_name": "营业成本审定表",
                                          "account_codes": ["6401~6499"]})
    if adj5_new:
        changes.append(f"F5: 新建块 {F5_SHEET_ADJ}")
    added = _ensure_cells(adj5, F5_ADJ_CELLS)
    for ref in added:
        changes.append(f"F5/{F5_SHEET_ADJ}: 新增 {ref}")

    # ─── F2 纠错 ──────────────────────────────────────────────────────────────
    changes.extend(_fix_f2_descriptions(doc))
    changes.extend(_fix_f2_range(doc))
    changes.extend(_merge_f5_old_block(doc))
    changes.extend(_fix_f5_pl_period(doc))

    # ─── F 类虚构 AUX 清理 ───────────────────────────────────────────────────
    changes.extend(_fix_aux_fabrications(doc))

    return changes


def validate(doc: dict[str, Any]) -> list[str]:
    """校验所有改动已到位（用于 --check）。返回欠账列表（空=通过）。"""
    issues: list[str] = []

    # F3 审定表
    f3_adj = _find_block(doc, "F3", F3_SHEET_ADJ)
    if f3_adj is None:
        issues.append(f"F3: 缺块 {F3_SHEET_ADJ}")
    else:
        refs = _cell_refs_in_block(f3_adj)
        for cell in F3_ADJ_CELLS:
            if cell["cell_ref"] not in refs:
                issues.append(f"F3/{F3_SHEET_ADJ}: 缺 {cell['cell_ref']}")

    # F3 明细
    f3_det = _find_block(doc, "F3", F3_SHEET_DETAIL)
    if f3_det is None:
        issues.append(f"F3: 缺块 {F3_SHEET_DETAIL}")
    else:
        refs = _cell_refs_in_block(f3_det)
        for cell in F3_DETAIL_BLOCK["cells"]:
            if cell["cell_ref"] not in refs:
                issues.append(f"F3/{F3_SHEET_DETAIL}: 缺 {cell['cell_ref']}")
        # 禁 WP 成环
        for cell in _cells(f3_det):
            if "WP(" in cell.get("formula", ""):
                issues.append(f"F3/{F3_SHEET_DETAIL}: 明细表禁 WP() → {cell.get('cell_ref','')}")

    # F3 披露两块
    for sheet, block_def in [(F3_SHEET_DISC_LISTED, F3_DISC_LISTED_BLOCK),
                             (F3_SHEET_DISC_SOE, F3_DISC_SOE_BLOCK)]:
        blk = _find_block(doc, "F3", sheet)
        if blk is None:
            issues.append(f"F3: 缺块 {sheet}")
        else:
            refs = _cell_refs_in_block(blk)
            for cell in block_def["cells"]:
                if cell["cell_ref"] not in refs:
                    issues.append(f"F3/{sheet}: 缺 {cell['cell_ref']}")

    # F4 同理
    f4_adj = _find_block(doc, "F4", F4_SHEET_ADJ)
    if f4_adj is None:
        issues.append(f"F4: 缺块 {F4_SHEET_ADJ}")
    else:
        refs = _cell_refs_in_block(f4_adj)
        for cell in F4_ADJ_CELLS:
            if cell["cell_ref"] not in refs:
                issues.append(f"F4/{F4_SHEET_ADJ}: 缺 {cell['cell_ref']}")

    f4_det = _find_block(doc, "F4", F4_SHEET_DETAIL)
    if f4_det is None:
        issues.append(f"F4: 缺块 {F4_SHEET_DETAIL}")
    else:
        refs = _cell_refs_in_block(f4_det)
        for cell in F4_DETAIL_BLOCK["cells"]:
            if cell["cell_ref"] not in refs:
                issues.append(f"F4/{F4_SHEET_DETAIL}: 缺 {cell['cell_ref']}")
        for cell in _cells(f4_det):
            if "WP(" in cell.get("formula", ""):
                issues.append(f"F4/{F4_SHEET_DETAIL}: 明细表禁 WP() → {cell.get('cell_ref','')}")

    for sheet, block_def in [(F4_SHEET_DISC_LISTED, F4_DISC_LISTED_BLOCK),
                             (F4_SHEET_DISC_SOE, F4_DISC_SOE_BLOCK)]:
        blk = _find_block(doc, "F4", sheet)
        if blk is None:
            issues.append(f"F4: 缺块 {sheet}")
        else:
            refs = _cell_refs_in_block(blk)
            for cell in block_def["cells"]:
                if cell["cell_ref"] not in refs:
                    issues.append(f"F4/{sheet}: 缺 {cell['cell_ref']}")

    # F5
    f5_adj = _find_block(doc, "F5", F5_SHEET_ADJ)
    if f5_adj is None:
        issues.append(f"F5: 缺块 {F5_SHEET_ADJ}")
    else:
        refs = _cell_refs_in_block(f5_adj)
        for cell in F5_ADJ_CELLS:
            if cell["cell_ref"] not in refs:
                issues.append(f"F5/{F5_SHEET_ADJ}: 缺 {cell['cell_ref']}")

    # F5 旧块不应存在
    old_f5 = _find_block(doc, "F5", F5_OLD_SHEET)
    if old_f5 is not None:
        issues.append(f"F5: 旧块 '{F5_OLD_SHEET}' 未合并删除")

    # 损益类禁 '期初余额'/'期末余额'（所有 F5 块）
    for b in _blocks(doc):
        if b.get("wp_code") != "F5":
            continue
        for cell in _cells(b):
            formula = cell.get("formula", "")
            if "'期初余额'" in formula or "'期末余额'" in formula:
                issues.append(
                    f"F5/{b.get('sheet','')}: 损益类禁期初/期末余额 → {cell.get('cell_ref','')}: {formula[:60]}"
                )

    # 虚构 AUX 清理
    for b in _blocks(doc):
        wp_code = b.get("wp_code", "")
        if not wp_code.startswith("F"):
            continue
        for cell in _cells(b):
            formula = cell.get("formula", "")
            for fab in F34_AUX_FABRICATIONS:
                if f"'{fab}'" in formula:
                    issues.append(
                        f"{wp_code}/{b.get('sheet','')}: 虚构AUX维度 '{fab}' 残留 → {cell.get('cell_ref','')}"
                    )
            for fab in F2_AUX_FABRICATIONS:
                if f"'{fab}'" in formula:
                    issues.append(
                        f"{wp_code}/{b.get('sheet','')}: 虚构AUX维度 '{fab}' 残留 → {cell.get('cell_ref','')}"
                    )

    # F2 口径检查
    f2_adj = _find_block(doc, "F2", "审定表F2-1")
    if f2_adj:
        for cell in _cells(f2_adj):
            formula = cell.get("formula", "")
            if "1401~1461" in formula:
                issues.append(
                    f"F2/审定表F2-1: 口径 1401~1461 未修正 → {cell.get('cell_ref','')}"
                )

    issues.extend(_validate_fc11_invariants(doc))
    return issues


# ═════════════════════ FC-11 不变量校验（本 spec Task 17 新增）═════════════════
#
# 🔴 改造前 `--check` 兼容读 `items`，于是「脚本报 0 项欠账」与「运行时一条公式都不生效」
#    可以同时成立。下面三条不变量把这个矛盾变成不可能：
#      ① 任何 F 循环块都不得有 `items` 键        —— 防回退（工具链根因已修，数据也不许残留）
#      ② 任何 F 循环块的 `cells` 必须非空        —— 空骨架块等于没配，不许静默通过
#      ③ 块的 `sheet` 必须是源 xlsx 的可见 tab 名 —— 全角/错字 sheet 名解析不到
#
# 🔴 作用域限定 F 循环：全库另有 4 个 `cells` 与 `items` 都空的块（`[171]` J1`明细表J1-2␠` /
#    `[179]` K1`明细表K1-2` / `[185]` K3`明细表K3-2` / `[195]` M2`明细表（非上市公司）M2-2`），
#    属其他循环的遗留欠账，不在本 spec 范围。把它们一并打红会让本校验永远无法转绿 ⇒
#    如实登记为顺带发现（见 evidence/task0-prerequisites.md §9），移交对应循环 spec。


def _visible_tabs_by_wp_code() -> dict[str, set[str]]:
    """读 `backend/wp_templates/**/{code} *.xlsx` 的可见 sheet tab 名（openpyxl 直读）。

    与 `backend/tests/test_f_cycle_formula_presets.py::f345_xlsx_tabs` 同源同法，
    差别只在这里覆盖全部 F 循环科目而非固定三家。openpyxl 缺失时返回空 dict
    （校验③自动跳过，不伪造通过也不误红）。
    """
    try:
        from openpyxl import load_workbook
    except ImportError:  # pragma: no cover - 环境缺 openpyxl 时降级
        return {}

    template_dir = _BACKEND / "wp_templates"
    tabs: dict[str, set[str]] = {}
    for path in sorted(template_dir.glob("**/*.xlsx")):
        if "~$" in path.name:
            continue
        code = path.stem.split(" ")[0]
        if not code.startswith("F"):
            continue
        try:
            wb = load_workbook(path, read_only=True, data_only=True)
        except Exception:  # pragma: no cover - 损坏模板不该让校验崩
            continue
        try:
            names = {sn for sn in wb.sheetnames if wb[sn].sheet_state == "visible"}
        finally:
            wb.close()
        tabs.setdefault(code, set()).update(names)
    return tabs


#: 🔴 跨 spec 已知欠账（**如实登记、不静默跳过**）—— `(wp_code, sheet)` → 移交说明。
#:
#: 这些块命中上面的不变量，但归属其他 spec 的范围。把它们计入 `issues` 会让本脚本的
#: `--check` 永远无法转绿（于是下一个人就会去放宽校验，FC-11 的老路）；直接从循环里
#: `continue` 掉又会让它们彻底消失（静默跳过，同样是 FC-11 的老路）。
#: 折中：单独归入 `[KNOWN-DEBT]` 段落**每次 `--check` 都打印**，但不判红。
#: 对应 spec 接手时删掉这里的条目，它就自动变成必须修的红项。
_KNOWN_CROSS_SPEC_DEBTS: dict[tuple[str, str], str] = {
    ("F0", "审定表F0-1"): (
        "sheet 名不存在于 F0 模板（真名 `函证结果汇总表F0-1`）⇒ 该块的 2 条 TB 公式"
        "运行时解析不到 sheet。归 F0/函证循环 spec；本批 F3/F4/F5 不越界改"
    ),
}


def _validate_fc11_invariants(doc: dict[str, Any]) -> list[str]:
    issues: list[str] = []
    known: list[str] = []
    tabs = _visible_tabs_by_wp_code()

    for idx, block in enumerate(_blocks(doc)):
        wp_code = str(block.get("wp_code") or "")
        if not wp_code.startswith("F"):
            continue
        sheet = str(block.get("sheet") or "")
        label = f"{wp_code}/{sheet}(块[{idx}])"

        debt_note = _KNOWN_CROSS_SPEC_DEBTS.get((wp_code, sheet))
        if debt_note is not None:
            known.append(f"{label}: {debt_note}")
            continue

        # ① 不得残留 items 键
        if "items" in block:
            issues.append(
                f"{label}: 残留 `items` 键（{len(block.get('items') or [])} 条）—— "
                "运行时四个消费方只读 `cells`，该键是死配置（FC-11）"
            )

        # ② cells 必须非空
        if not _cells(block):
            issues.append(
                f"{label}: `cells` 为空 —— 空骨架块在公式管理页看不到、不预填、不建边（FC-11）"
            )

        # ③ sheet 名必须是源 xlsx 可见 tab
        if sheet and wp_code in tabs and sheet not in tabs[wp_code]:
            issues.append(
                f"{label}: sheet 名不在源 xlsx 可见 tab 集合中 —— "
                f"可用 tab: {sorted(tabs[wp_code])}"
            )

    if known:
        print(f"[KNOWN-DEBT] {len(known)} 项跨 spec 已知欠账（不判红，见 _KNOWN_CROSS_SPEC_DEBTS）:")
        for item in known:
            print(f"  · {item}")

    return issues


def main() -> int:
    parser = argparse.ArgumentParser(description="补齐 F3/F4/F5 公式预设 + 纠错 F2/F5")
    parser.add_argument("--dry-run", action="store_true", help="只打印变更不写盘")
    parser.add_argument("--check", action="store_true", help="校验模式：打印欠账并返回 exit code")
    args = parser.parse_args()

    with open(MAPPING_PATH, encoding="utf-8") as f:
        doc = json.load(f)

    if args.check:
        issues = validate(doc)
        if issues:
            print(f"[FAIL] {len(issues)} 项欠账:")
            for iss in issues:
                print(f"  - {iss}")
            return 1
        print("[OK] F3/F4/F5 公式预设全部到位，0 项欠账")
        return 0

    changes = apply(doc)
    if not changes:
        print("[OK] 无需修改（幂等）")
        return 0

    print(f"变更 {len(changes)} 项:")
    for c in changes:
        print(f"  - {c}")

    if args.dry_run:
        print("\n(--dry-run 模式，未写盘)")
        return 0

    # 写盘
    with open(MAPPING_PATH, "w", encoding="utf-8") as f:
        json.dump(doc, f, ensure_ascii=False, indent=2)
    print(f"\n[OK] 已写入 {MAPPING_PATH}")

    # round-trip 自检
    with open(MAPPING_PATH, encoding="utf-8") as f:
        reloaded = json.load(f)
    issues = validate(reloaded)
    if issues:
        print(f"[WARN] round-trip 自检失败（{len(issues)} 项欠账），请检查！")
        return 2

    print("[OK] round-trip 自检通过")
    return 0


if __name__ == "__main__":
    sys.exit(main())

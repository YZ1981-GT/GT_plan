# -*- coding: utf-8 -*-
"""净化 N4 税金及附加 权威模板的外部关系（过首版发布 OOXML `external_relationships` 门）。

spec: `n-cycle-sync-foundation-and-first-canary` · N4 canary · Task 7a

范式：照 `sanitize_i_cycle_template_external_links.py` / `sanitize_l4_template_external_links.py`
—— `backend/wp_templates/` 运行时只读，模板真要升级必须显式净化 + 留 `.preclean.bak`
作门负例 + 重算 sha256 + 重生成契约 + 重发布。**复用 I 循环的 `sanitize_bytes` 内核**
（它已覆盖 externalLinks / embeddings-OLE / 外部 hyperlink / 外部引用 defined name 四类）。

逐册现查（2026-10-01，zip 内逐部件扫描；`_n4p_ooxml.py` 探针现算）：

| 项 | 现算值 | 说明 |
|---|---|---|
| 外链部件 | **4**（`externalLink{1,2}.xml` + 各自 `_rels`） | 旧作者本机引用，断链 |
| 含 `[n]` 公式 | **5** | 全在**隐藏**册 `税金及附加审计程序表O2A（原底稿）`，引用 `[1]底稿目录!`/`[2]底稿目录!` |
| 含 `[n]` defined name | 0 | |
| 外部 hyperlink | 0 | |
| OLE 嵌入对象 | 0 | `vmlDrawing{1,2}.vml` 是批注/控件 VML，**不触发门**，保留 |
| `.bin` | 9 个 `printerSettings`（打印设置，非 OLE） | 门不拒 |

处置（与 I/J/D 系同口径）：删外链部件 + 相关 Relationship/Override/externalReferences；
含 `[n]` 的 5 个公式去 `<f>` 保 `<v>` 缓存值（它们落在隐藏「原底稿」册，OO 打开本就取不到
外部工作簿 ⇒ 转缓存值是降级不是破坏）。**受管 sheet `税金及附加明细表N4-2` 逐格 0 diff**
（探针实证 77→77 格、merge 不变）；派生表 `税金及附加审定表N4-1` 不受影响。

用法（仓库根）：
  & .venv/Scripts/python.exe backend/scripts/fix/sanitize_n4_template_external_links.py          # 预演
  & .venv/Scripts/python.exe backend/scripts/fix/sanitize_n4_template_external_links.py --apply  # 写盘 + .bak

🔴 判成败（不看退出码）：①净化后 OOXML 门 PASS 且 `.preclean.bak` 仍 REJECT ②受管 sheet 逐格
0 diff + merge 不变 ③各项计数与上表**逐值相等**（多一处少一处都拒绝写盘）。
净化后 sha256 = `2005eada32506e9e2b1b6f68c704ca4f6626b8a78e9e2a3a221ca00602cad638`。
"""
from __future__ import annotations

import hashlib
import io
import shutil
import sys
import tempfile
from dataclasses import dataclass
from pathlib import Path

_REPO = Path(__file__).resolve().parents[3]
sys.path.insert(0, str(_REPO / "backend"))

from scripts.fix.sanitize_i_cycle_template_external_links import sanitize_bytes  # noqa: E402

_TPL_DIR = _REPO / "backend" / "wp_templates" / "N"


@dataclass(frozen=True)
class _Target:
    file_name: str
    managed_sheet: str
    expect_parts: int
    expect_formulas: int
    expect_defined_names: int
    expect_hyperlinks: int
    expect_ole_objects: int


#: 计数口径同 I：`expect_parts` = 外链部件 + 其 `_rels`（一条外链 = 2 个部件）。
TARGET = _Target(
    file_name="N4 税金及附加.xlsx",
    managed_sheet="税金及附加明细表N4-2",
    expect_parts=4,
    expect_formulas=5,
    expect_defined_names=0,
    expect_hyperlinks=0,
    expect_ole_objects=0,
)


def _snapshot(data: bytes, sheet: str):
    import openpyxl

    wb = openpyxl.load_workbook(io.BytesIO(data), data_only=False)
    ws = wb[sheet]

    def _val(v: object) -> str:
        text = getattr(v, "text", None)
        return f"{{array {getattr(v, 'ref', '')}}}{text}" if text is not None else str(v)

    snap = {
        c.coordinate: _val(c.value)
        for row in ws.iter_rows(min_row=1, max_row=ws.max_row, max_col=ws.max_column)
        for c in row
        if c.value is not None
    }
    merged = sorted(str(m) for m in ws.merged_cells.ranges)
    wb.close()
    return snap, merged


def _gate(data: bytes):
    from app.services.workpaper_sync.ooxml_security import validate_ooxml_artifact

    with tempfile.TemporaryDirectory() as tmp:
        p = Path(tmp) / "probe.xlsx"
        p.write_bytes(data)
        try:
            validate_ooxml_artifact(p, document_type="xlsx")
            return True, "PASS"
        except Exception as exc:  # noqa: BLE001 - 门负例要拿到拒绝原因
            return False, f"REJECT {getattr(exc, 'gate', type(exc).__name__)}"


def run_one(t: _Target, *, apply: bool) -> bool:
    path = _TPL_DIR / t.file_name
    bak = path.with_suffix(path.suffix + ".preclean.bak")
    src = bak.read_bytes() if bak.exists() else path.read_bytes()
    out, stats = sanitize_bytes(src)
    b_snap, b_merged = _snapshot(src, t.managed_sheet)
    a_snap, a_merged = _snapshot(out, t.managed_sheet)
    diffs = [k for k in sorted(set(b_snap) | set(a_snap)) if b_snap.get(k) != a_snap.get(k)]
    gate_before, gate_after = _gate(src), _gate(out)
    expected = {
        "dropped_parts": t.expect_parts,
        "neutralized_formulas": t.expect_formulas,
        "dropped_defined_names": t.expect_defined_names,
        "dropped_hyperlinks": t.expect_hyperlinks,
        "dropped_ole_objects": t.expect_ole_objects,
    }
    print(f"== {t.file_name}")
    print(f"   before sha256={hashlib.sha256(src).hexdigest()} size={len(src)} gate={gate_before[1]}")
    print(f"   after  sha256={hashlib.sha256(out).hexdigest()} size={len(out)} gate={gate_after[1]}")
    print(f"   stats={stats} expected={expected}")
    print(f"   managed {t.managed_sheet}: cells {len(b_snap)}->{len(a_snap)} diffs={diffs[:8]} merged_same={b_merged == a_merged}")
    ok = (not gate_before[0]) and gate_after[0] and not diffs and b_merged == a_merged and stats == expected
    if not ok:
        print("   [FAIL] 判据不满足 —— 不写盘")
        return False
    if apply:
        if not bak.exists():
            shutil.copy2(path, bak)
            print(f"   [apply] 备份门负例 -> {bak.name}")
        if path.read_bytes() != out:
            path.write_bytes(out)
            print("   [apply] 已净化写盘")
        else:
            print("   [apply] 已是净化态（幂等）")
    return True


def main() -> int:
    apply = "--apply" in sys.argv[1:]
    ok = run_one(TARGET, apply=apply)
    print("ALL OK" if ok else "SOME FAILED")
    return 0 if ok else 1


if __name__ == "__main__":
    raise SystemExit(main())

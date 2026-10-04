# -*- coding: utf-8 -*-
"""净化 I 循环 I2~I5 权威模板的外部关系（过首版发布 OOXML `external_relationships` 门）。

spec: `i-cycle-sync-foundation-and-first-canary` + `i1-i3-…` + `i2-i4-i5-…`
（2026-10-01 首版发布 `--check` 四册被 `ooxml_security_rejected` 挡下后）

范式：照 `sanitize_d{3..7}_template_external_links.py` / `sanitize_j1_template_external_links.py`
—— `backend/wp_templates/` 运行时只读，模板真要升级必须显式净化 + 留 `.preclean.bak`
作门负例 + 重算 sha256 + 重生成契约 + 重发布。I1 / I6 两册**无**外部关系，不动。

逐册现查（2026-10-01，zip 内逐部件扫描）：

| 册 | 外链部件 | `[n]` 公式 | `[n]` defined name | 外部 hyperlink |
|---|---|---|---|---|
| I2 开发支出 | 2（`I6 研发费用.xlsx`） | **14**（全在 `实质性分析I2-5`，`'[1]明细表I6-2'!…`） | 0 | 0 |
| I3 商誉 | 0 | 0 | 0 | **2**（`商誉减值测试I3-6` → gt-china.com.cn 文章） |
| I4 长期待摊费用 | 20（旧作者本机 `A:\` `E:\` `Y:\` 路径） | 0 | **107** | 0 |
| I5 其他非流动资产 | 20（同上） | 0 | **106** | 0 |

处置（与 J1 / D 系同口径）：
  ① 删 `xl/externalLinks/*` 部件与其 `_rels`；`[Content_Types].xml` 删对应 Override；
     `xl/_rels/workbook.xml.rels` 删对应 Relationship；`xl/workbook.xml` 删 `<externalReferences>`
  ② 含 `[n]` 的 defined name 删除（全部指向已断链外部工作簿，任何公式引用它们都只会得 `#REF!`）
  ③ 含 `[n]` 的公式格去 `<f>` 保 `<v>` 缓存值（可见内容不变）。🔴 I2-5 那 14 格是跨册引用
     I6 明细，在平台里本就取不到（OO 打开即刷新提示 + `#REF!`），转缓存值是降级不是破坏
  ④ sheet 级 `TargetMode="External"` 的 hyperlink：删 Relationship 与 `<hyperlink r:id=…>` 元素，
     **单元格文字保留**（I3-6 两处参考文章链接失去可点击性，内容不丢）

用法（仓库根）：
  & .venv/Scripts/python.exe backend/scripts/fix/sanitize_i_cycle_template_external_links.py          # 预演
  & .venv/Scripts/python.exe backend/scripts/fix/sanitize_i_cycle_template_external_links.py --apply  # 写盘 + .bak

🔴 判成败：①净化后 OOXML 门 PASS 且 `.preclean.bak` 仍 REJECT ②受管 sheet 逐格 0 diff + merge 不变
③各项计数与上表**逐值相等**（多一处少一处都拒绝写盘）。不看退出码。
"""
from __future__ import annotations

import hashlib
import io
import re
import shutil
import sys
import zipfile
from dataclasses import dataclass
from pathlib import Path

_REPO = Path(__file__).resolve().parents[3]
sys.path.insert(0, str(_REPO / "backend"))
_TPL_DIR = _REPO / "backend" / "wp_templates" / "I"


@dataclass(frozen=True)
class _Target:
    file_name: str
    managed_sheet: str
    expect_parts: int  # 外链部件 + 其 rels
    expect_formulas: int
    expect_defined_names: int
    expect_hyperlinks: int
    expect_ole_objects: int


#: 计数口径：`expect_parts` = 外链部件 + 其 `_rels`（一条外链 = 2 个部件）+ `xl/embeddings/*`。
TARGETS: tuple[_Target, ...] = (
    _Target("I2 开发支出.xlsx", "明细表I2-2", 2, 14, 0, 0, 0),
    #: 🔴 I3 去掉 2 个外部 hyperlink 后被**第二道门** `embedded_objects` 拒：17 个 OLE 嵌入对象
    #: （`可收回金额测试I3-7` 2 个 + `参考－商誉减值测试示例` 15 个，公式编辑器对象）
    _Target("I3 商誉.xlsx", "明细表I3-2", 17, 0, 0, 2, 17),
    _Target("I4 长期待摊费用.xlsx", "明细表I4-2", 20, 0, 107, 0, 0),
    _Target("I5 其他非流动资产.xlsx", "明细表I5-2", 20, 0, 106, 0, 0),
)

_EXT_LINK_PART_RE = re.compile(r"^xl/externalLinks/externalLink\d+\.xml$")
_EXT_LINK_RELS_RE = re.compile(r"^xl/externalLinks/_rels/externalLink\d+\.xml\.rels$")
_EXT_FORMULA_RE = re.compile(r"<f>[^<]*\[\d+\][^<]*</f>")
_EXT_REL_RE = re.compile(r'<Relationship\b[^>]*TargetMode="External"[^>]*/>')
_EMBED_PART_RE = re.compile(r"^xl/embeddings/")
_EMBED_REL_RE = re.compile(r'<Relationship\b[^>]*Target="[^"]*embeddings/[^"]*"[^>]*/>')
#: `<oleObjects>` 容器（Excel 2010+ 里每个对象包在 `mc:AlternateContent` 里，整块删最干净）。
#: 对象的 VML 预览形状与 `xl/media/` 图片保留 ⇒ 视觉上仍显示公式图，只是不可双击编辑。
_OLE_BLOCK_RE = re.compile(r"<oleObjects>.*?</oleObjects>", re.S)


def _strip_workbook(xml: str) -> tuple[str, int]:
    xml = re.sub(r"<externalReferences>.*?</externalReferences>", "", xml, flags=re.S)
    dropped = 0

    def _drop(m: re.Match) -> str:
        nonlocal dropped
        if re.search(r"\[\d+\]", m.group(0)):
            dropped += 1
            return ""
        return m.group(0)

    xml = re.sub(r"<definedName\b[^>]*>.*?</definedName>", _drop, xml, flags=re.S)
    # 删光后可能留下空 <definedNames></definedNames>（Excel 拒收空容器）
    xml = re.sub(r"<definedNames>\s*</definedNames>", "", xml)
    return xml, dropped


def sanitize_bytes(src: bytes) -> tuple[bytes, dict[str, int]]:
    zin = zipfile.ZipFile(io.BytesIO(src))
    stats = {"dropped_parts": 0, "neutralized_formulas": 0, "dropped_defined_names": 0,
             "dropped_hyperlinks": 0, "dropped_ole_objects": 0}
    # 先算每个 sheet 被删的外部 hyperlink rId（rels 与 sheet xml 要配对改）
    drop_rids: dict[str, set[str]] = {}
    for name in zin.namelist():
        m = re.match(r"^xl/worksheets/_rels/(sheet\d+\.xml)\.rels$", name)
        if not m:
            continue
        rels = zin.read(name).decode("utf-8")
        ids = {re.search(r'Id="([^"]+)"', r).group(1) for r in _EXT_REL_RE.findall(rels)}
        if ids:
            drop_rids[f"xl/worksheets/{m.group(1)}"] = ids

    out = io.BytesIO()
    zout = zipfile.ZipFile(out, "w", zipfile.ZIP_DEFLATED)
    for info in zin.infolist():
        name = info.filename
        if _EXT_LINK_PART_RE.match(name) or _EXT_LINK_RELS_RE.match(name):
            stats["dropped_parts"] += 1
            continue
        if _EMBED_PART_RE.match(name):
            stats["dropped_parts"] += 1
            continue
        data = zin.read(name)
        if name == "[Content_Types].xml":
            text = re.sub(
                r'<Override PartName="/xl/externalLinks/externalLink\d+\.xml"[^>]*/>', "",
                data.decode("utf-8"),
            )
            text = re.sub(r'<Override PartName="/xl/embeddings/[^"]*"[^>]*/>', "", text)
            data = text.encode("utf-8")
        elif name == "xl/_rels/workbook.xml.rels":
            data = re.sub(
                r'<Relationship [^>]*Target="externalLinks/externalLink\d+\.xml"[^>]*/>', "",
                data.decode("utf-8"),
            ).encode("utf-8")
        elif name == "xl/workbook.xml":
            text, n = _strip_workbook(data.decode("utf-8"))
            stats["dropped_defined_names"] += n
            data = text.encode("utf-8")
        elif re.match(r"^xl/worksheets/_rels/sheet\d+\.xml\.rels$", name):
            text = data.decode("utf-8")
            text = _EXT_REL_RE.sub("", text)
            text = _EMBED_REL_RE.sub("", text)
            data = text.encode("utf-8")
        elif name.startswith("xl/worksheets/sheet") and name.endswith(".xml"):
            text = data.decode("utf-8")
            hits = len(_EXT_FORMULA_RE.findall(text))
            if hits:
                stats["neutralized_formulas"] += hits
                text = _EXT_FORMULA_RE.sub("", text)
            for rid in sorted(drop_rids.get(name, ())):
                text, n = re.subn(r'<hyperlink\b[^>]*r:id="' + re.escape(rid) + r'"[^>]*/>', "", text)
                stats["dropped_hyperlinks"] += n
            text = re.sub(r"<hyperlinks>\s*</hyperlinks>", "", text)
            for block in _OLE_BLOCK_RE.findall(text):
                # 每个对象在 Choice 与 Fallback 各出现一次 `<oleObject `，按 r:id 去重计数
                stats["dropped_ole_objects"] += len(set(re.findall(r'<oleObject\b[^>]*r:id="([^"]+)"', block)))
            text = _OLE_BLOCK_RE.sub("", text)
            data = text.encode("utf-8")
        zout.writestr(info, data)
    zout.close()
    zin.close()
    return out.getvalue(), stats


def _snapshot(data: bytes, sheet: str) -> tuple[dict[str, str], list[str]]:
    import openpyxl

    wb = openpyxl.load_workbook(io.BytesIO(data), data_only=False)
    ws = wb[sheet]
    def _val(v: object) -> str:
        # 🔴 ArrayFormula 的 str() 是对象地址 —— 按它比会把未改动的数组公式全报成 diff
        text = getattr(v, "text", None)
        return f"{{array {getattr(v, 'ref', '')}}}{text}" if text is not None else str(v)

    snap = {
        c.coordinate: _val(c.value)
        for row in ws.iter_rows(min_row=1, max_row=ws.max_row, max_col=ws.max_column)
        for c in row
        if c.value is not None
    }
    merged = sorted(str(m) for m in ws.merged_cells.ranges)
    wb.close()
    return snap, merged


def _gate(data: bytes) -> tuple[bool, str]:
    """真跑生产 OOXML 门（与首版发布同一入口）。"""
    import tempfile

    from app.services.workpaper_sync.ooxml_security import validate_ooxml_artifact

    with tempfile.TemporaryDirectory() as tmp:
        p = Path(tmp) / "probe.xlsx"
        p.write_bytes(data)
        try:
            validate_ooxml_artifact(p, document_type="xlsx")
            return True, "PASS"
        except Exception as exc:  # noqa: BLE001 - 门负例要拿到拒绝原因
            return False, f"REJECT {getattr(exc, 'gate', type(exc).__name__)}"


def run_one(t: _Target, *, apply: bool) -> bool:
    path = _TPL_DIR / t.file_name
    bak = path.with_suffix(path.suffix + ".preclean.bak")
    src = bak.read_bytes() if bak.exists() else path.read_bytes()
    out, stats = sanitize_bytes(src)
    b_snap, b_merged = _snapshot(src, t.managed_sheet)
    a_snap, a_merged = _snapshot(out, t.managed_sheet)
    diffs = [k for k in sorted(set(b_snap) | set(a_snap)) if b_snap.get(k) != a_snap.get(k)]
    gate_before, gate_after = _gate(src), _gate(out)
    expected = {"dropped_parts": t.expect_parts, "neutralized_formulas": t.expect_formulas,
                "dropped_defined_names": t.expect_defined_names,
                "dropped_hyperlinks": t.expect_hyperlinks,
                "dropped_ole_objects": t.expect_ole_objects}
    print(f"== {t.file_name}")
    print(f"   before sha256={hashlib.sha256(src).hexdigest()} size={len(src)} gate={gate_before[1]}")
    print(f"   after  sha256={hashlib.sha256(out).hexdigest()} size={len(out)} gate={gate_after[1]}")
    print(f"   stats={stats} expected={expected}")
    print(f"   managed {t.managed_sheet}: cells {len(b_snap)}->{len(a_snap)} diffs={diffs[:8]} merged_same={b_merged == a_merged}")
    ok = (not gate_before[0]) and gate_after[0] and not diffs and b_merged == a_merged and stats == expected
    if not ok:
        print("   [FAIL] 判据不满足 —— 不写盘")
        return False
    if apply:
        if not bak.exists():
            shutil.copy2(path, bak)
            print(f"   [apply] 备份门负例 -> {bak.name}")
        if path.read_bytes() != out:
            path.write_bytes(out)
            print("   [apply] 已净化写盘")
        else:
            print("   [apply] 已是净化态（幂等）")
    return True


def main() -> int:
    apply = "--apply" in sys.argv[1:]
    ok = all([run_one(t, apply=apply) for t in TARGETS])
    print("ALL OK" if ok else "SOME FAILED")
    return 0 if ok else 1


if __name__ == "__main__":
    raise SystemExit(main())

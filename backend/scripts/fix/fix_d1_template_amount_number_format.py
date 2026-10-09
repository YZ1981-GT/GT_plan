# -*- coding: utf-8 -*-
"""修 D1 权威模板里「契约声明 amount 但单元格是日期格式」的格子（zip 级精修）。

spec: d1-sync-row-table-engine-and-d1-coverage · Tasks 25~29 的整册门（V5 节裁决）

═══ 修的是什么 ═══

`应收票据备查簿核对D1-7` 的 **N 列「贴现息」** 在两张行表（银行承兑 / 商业承兑）里
都带 `mm-dd-yy` 数字格式，而契约声明它是 `amount`。后果有两层：

1. **数据损坏**：materialize 写进金额 `0`，openpyxl 按日期格式反读回
   `datetime.time(0, 0)` ⇒ G1 roundtrip 等值门以
   `ValueNormalizationError: amount 字段收到 time` 拒收整册 materialize。
2. **审计师肉眼可见**：Excel 里「贴现息 5000」会显示成一个日期。

🔴 判据是**客观**的，不是「格式看起来怪」：全 D1 的 12 张受管 sheet / 18 张行表逐字段
扫过，「契约声明 amount 且单元格 number_format 是日期/时间形态」**只有这 2 处**。
文本列上的怪格式（名称列挂欧元货币、状态列挂日期）**不在修复范围** —— Excel 的数字
格式只作用于数值，文本单元格照原样显示，那些既不损坏数据也不可见。

═══ 为什么走 zip 级精修而不是 openpyxl 重存 ═══

实测 openpyxl 读入再保存本模板：138008 → 121363 字节，且会**删掉** `printerSettings`
/ 批注 / `calcChain`、重命名 Table 部件、改写 3 个 definedName。对权威审计底稿来说
那是不可接受的附带损伤。本脚本只动两个 zip 部件（`xl/styles.xml` 与目标 sheet 的
XML），其余部件逐字节原样搬运。

═══ 为什么新建 cellXfs 而不是直接套用 H 列的 style ═══

H 列（票据金额）用的是 `s="97"`，格式正确 —— 但直接把 N 列改成 97 会连 H 的字体、
边框、对齐、填充一起带过来，N 列在受管区边缘的边框会走形。本脚本**复制 N 列现有的
xf（`s="96"`）并只替换它的 `numFmtId`**，做到「除数字格式外一切不变」。

用法：`--check` 只报告不写盘；`--apply` 写盘。
写盘后 `TEMPLATE_SHA256` 会变，必须按顺序重生成：
  1. `backend/scripts/gen/generate_phase5_d1_contract.py --apply`
  2. `backend/scripts/fix/fix_d1_republish_representation.py --apply`
"""
from __future__ import annotations

import argparse
import hashlib
import io
import json
import re
import shutil
import sys
import zipfile
from pathlib import Path
from typing import Any

_BACKEND = Path(__file__).resolve().parents[2]
if str(_BACKEND) not in sys.path:
    sys.path.insert(0, str(_BACKEND))

for _stream in (sys.stdout, sys.stderr):
    try:
        _stream.reconfigure(encoding="utf-8")  # type: ignore[attr-defined]
    except Exception:  # pragma: no cover
        pass

#: 受修的 sheet 与列。**不写死行号** —— 行号从 provider 的行表 spec 现取。
TARGET_SHEET = "应收票据备查簿核对D1-7"
TARGET_COLUMN = "N"
#: 该列在两张行表里的字段名（契约声明 amount 的那个）。
TARGET_FIELD = "discount_interest"
#: 取正确数字格式的参照列（同 sheet 的「票据金额」，契约同样声明 amount）。
REFERENCE_COLUMN = "H"


def _col_index(letter: str) -> int:
    n = 0
    for ch in letter.upper():
        n = n * 26 + (ord(ch) - 64)
    return n


def _target_rows() -> list[int]:
    """从 provider 的行表 spec 现取受管数据行（不写死）。"""
    from app.services.workpaper_sync.phase5_d1_expansion import managed_row_table_specs

    rows: list[int] = []
    for spec in managed_row_table_specs():
        keys = {
            str(getattr(c, "field_key", "") or getattr(c, "key", ""))
            for c in (getattr(spec, "columns", None) or ())
        }
        table_key = str(getattr(spec, "table_key", ""))
        if not table_key.startswith("memo_"):
            continue
        first = int(getattr(spec, "first_data_row", 0) or 0)
        last = int(getattr(spec, "last_data_row", 0) or 0)
        if first and last:
            rows.extend(range(first, last + 1))
        del keys
    return sorted(set(rows))


def _sheet_part(zf: zipfile.ZipFile, sheet_name: str) -> str:
    wbxml = zf.read("xl/workbook.xml").decode("utf-8")
    rels = zf.read("xl/_rels/workbook.xml.rels").decode("utf-8")
    m = re.search(
        r'<sheet[^>]*name="%s"[^>]*r:id="(rId\d+)"' % re.escape(sheet_name), wbxml
    )
    if not m:
        raise SystemExit(f"模板里找不到 sheet {sheet_name!r}")
    m2 = re.search(r'Id="%s"[^>]*Target="([^"]+)"' % re.escape(m.group(1)), rels)
    if not m2:
        raise SystemExit(f"找不到 {m.group(1)} 的 Target")
    return "xl/" + m2.group(1).lstrip("/")


def _cellxfs(styles: str) -> tuple[int, int, list[str]]:
    """返回 `(cellXfs 内容起点, 终点, xf 元素列表)`。

    🔴 不能用 `<xf\\b.*?(?:/>|</xf>)` 这种非贪婪正则：`<xf>` 可以带
    `<alignment/>` / `<protection/>` 子元素，非贪婪会在**子元素的** `/>` 处截断，
    切出 `<xf ...><alignment/>` 这种半个元素，复制出去就是
    `XMLSyntaxError: Opening and ending tag mismatch`。按开闭标签顺序扫描。
    """
    m = re.search(r"<cellXfs\b[^>]*>(.*?)</cellXfs>", styles, re.S)
    if not m:
        raise SystemExit("styles.xml 里没有 cellXfs")
    inner = m.group(1)
    xfs: list[str] = []
    pos = 0
    while True:
        start = inner.find("<xf", pos)
        if start < 0:
            break
        gt = inner.find(">", start)
        if gt < 0:
            raise SystemExit("cellXfs 里有未闭合的 <xf 开标签")
        if inner[gt - 1] == "/":          # 自闭合 <xf ... />
            xfs.append(inner[start : gt + 1])
            pos = gt + 1
            continue
        close = inner.find("</xf>", gt)   # 带子元素 <xf ...>…</xf>
        if close < 0:
            raise SystemExit("cellXfs 里有 <xf> 缺 </xf>")
        xfs.append(inner[start : close + len("</xf>")])
        pos = close + len("</xf>")
    return m.start(1), m.end(1), xfs


def _num_fmt_id(xf: str) -> str:
    m = re.search(r'numFmtId="(\d+)"', xf)
    return m.group(1) if m else "0"


def _style_of(sheet_xml: str, coord: str) -> str | None:
    m = re.search(r'<c r="%s"(?=[ />])[^>]*\bs="(\d+)"' % coord, sheet_xml)
    return m.group(1) if m else None


def _number_formats(data: bytes) -> dict[str, str]:
    """整簿 `(sheet!coord) -> number_format` 快照，用于证明改动面最小。"""
    import openpyxl

    wb = openpyxl.load_workbook(io.BytesIO(data))
    out: dict[str, str] = {}
    for ws in wb.worksheets:
        for row in ws.iter_rows():
            for cell in row:
                out[f"{ws.title}!{cell.coordinate}"] = str(cell.number_format)
    wb.close()
    return out


def run(*, apply: bool) -> dict[str, Any]:
    from app.services.workpaper_sync import phase5_d1_notes_receivable as P

    template_path = (
        _BACKEND / "wp_templates" / Path(P.TEMPLATE_RELATIVE_PATH)
    )
    if not template_path.exists():
        raise SystemExit(f"模板不存在: {template_path}")
    original = template_path.read_bytes()
    report: dict[str, Any] = {
        "template": str(template_path.relative_to(_BACKEND)),
        "sha256_before": hashlib.sha256(original).hexdigest(),
        "mode": "apply" if apply else "check",
    }

    rows = _target_rows()
    if not rows:
        raise SystemExit("受管行区间算出为空 —— 空分母，拒绝写盘")
    report["target_rows"] = rows

    zf = zipfile.ZipFile(io.BytesIO(original))
    part = _sheet_part(zf, TARGET_SHEET)
    report["sheet_part"] = part
    sheet_xml = zf.read(part).decode("utf-8")
    styles_xml = zf.read("xl/styles.xml").decode("utf-8")

    # ── 目标格与参照格的 style id ──
    target_styles = {
        r: _style_of(sheet_xml, f"{TARGET_COLUMN}{r}") for r in rows
    }
    missing = sorted(r for r, s in target_styles.items() if s is None)
    if missing:
        raise SystemExit(
            f"{TARGET_COLUMN} 列这些行在 XML 里没有显式 s= : {missing} —— "
            "本脚本只改显式带 style 的格，遇到隐式格应先确认格式来源（<col style=>）"
        )
    distinct = sorted({s for s in target_styles.values() if s})
    report["target_style_ids"] = distinct
    ref_style = _style_of(sheet_xml, f"{REFERENCE_COLUMN}{rows[0]}")
    report["reference_style_id"] = ref_style
    if ref_style is None:
        raise SystemExit(f"参照格 {REFERENCE_COLUMN}{rows[0]} 没有显式 s=")

    _start, end, xfs = _cellxfs(styles_xml)
    ref_fmt_id = _num_fmt_id(xfs[int(ref_style)])
    report["reference_num_fmt_id"] = ref_fmt_id

    # ── 为每个待改 style 造一个「只换 numFmtId」的新 xf ──
    new_xfs: list[str] = []
    remap: dict[str, str] = {}
    for sid in distinct:
        src = xfs[int(sid)]
        if _num_fmt_id(src) == ref_fmt_id:
            continue  # 已是目标格式
        patched = re.sub(r'numFmtId="\d+"', f'numFmtId="{ref_fmt_id}"', src, count=1)
        if 'applyNumberFormat="1"' not in patched:
            patched = re.sub(r"<xf\b", '<xf applyNumberFormat="1"', patched, count=1)
        remap[sid] = str(len(xfs) + len(new_xfs))
        new_xfs.append(patched)
    report["remap"] = remap
    report["new_xf_count"] = len(new_xfs)
    if not remap:
        report["status"] = "already_correct"
        print("[=] 目标格已是正确数字格式，无需改动")
        return report

    # ── 改 sheet XML：只改目标列目标行的 s= ──
    changed = 0

    def _patch_cell(match: re.Match[str]) -> str:
        nonlocal changed
        tag = match.group(0)
        sm = re.search(r'\bs="(\d+)"', tag)
        if not sm or sm.group(1) not in remap:
            return tag
        changed += 1
        return tag[: sm.start(1)] + remap[sm.group(1)] + tag[sm.end(1) :]

    for r in rows:
        sheet_xml, n = re.subn(
            r'<c r="%s%d"(?=[ />])[^>]*?(?:/>|>)' % (TARGET_COLUMN, r),
            _patch_cell,
            sheet_xml,
            count=1,
        )
        del n
    report["cells_changed"] = changed
    if changed != len(rows):
        raise SystemExit(
            f"预期改 {len(rows)} 个格，实际改了 {changed} 个 —— 拒绝写盘"
        )

    # ── 改 styles.xml：追加新 xf 并更新 count ──
    new_styles = styles_xml[:end] + "".join(new_xfs) + styles_xml[end:]
    new_styles = re.sub(
        r'(<cellXfs\b[^>]*\bcount=")(\d+)(")',
        lambda m: m.group(1) + str(len(xfs) + len(new_xfs)) + m.group(3),
        new_styles,
        count=1,
    )

    # ── 重打包：只替换两个部件，其余逐字节搬运 ──
    buf = io.BytesIO()
    with zipfile.ZipFile(buf, "w", zipfile.ZIP_DEFLATED) as out:
        for info in zf.infolist():
            if info.filename == part:
                payload = sheet_xml.encode("utf-8")
            elif info.filename == "xl/styles.xml":
                payload = new_styles.encode("utf-8")
            else:
                payload = zf.read(info.filename)
            # 🔴 不能把 `infolist()` 拿到的 ZipInfo 直接交给 writestr —— 它带着**原**
            #    CRC 与 compress_size，配上新字节会写出 `BadZipFile: Bad CRC-32`。
            #    新建 ZipInfo 只搬运需要保真的元数据（名字/时间戳/压缩方式/权限位）。
            fresh = zipfile.ZipInfo(info.filename, date_time=info.date_time)
            fresh.compress_type = info.compress_type
            fresh.external_attr = info.external_attr
            fresh.internal_attr = info.internal_attr
            fresh.create_system = info.create_system
            out.writestr(fresh, payload)
    patched_bytes = buf.getvalue()
    report["sha256_after"] = hashlib.sha256(patched_bytes).hexdigest()
    report["size_before"] = len(original)
    report["size_after"] = len(patched_bytes)

    # ── 前置校验①：除目标格外，全簿 number_format 一个都不许变 ──
    before_fmts = _number_formats(original)
    after_fmts = _number_formats(patched_bytes)
    expected = {f"{TARGET_SHEET}!{TARGET_COLUMN}{r}" for r in rows}
    drifted = {
        k for k in set(before_fmts) | set(after_fmts)
        if before_fmts.get(k) != after_fmts.get(k)
    }
    unexpected = sorted(drifted - expected)
    if unexpected:
        raise SystemExit(
            f"改动面超出预期：{len(unexpected)} 个非目标格的数字格式也变了，"
            f"前 8 个 = {unexpected[:8]} —— 拒绝写盘"
        )
    not_changed = sorted(expected - drifted)
    if not_changed:
        raise SystemExit(f"这些目标格的数字格式没变：{not_changed} —— 拒绝写盘")
    report["format_changed_cells"] = sorted(expected)
    report["reference_format"] = after_fmts.get(f"{TARGET_SHEET}!{REFERENCE_COLUMN}{rows[0]}")
    report["new_format"] = after_fmts.get(f"{TARGET_SHEET}!{TARGET_COLUMN}{rows[0]}")

    # ── 前置校验②：zip 部件集合不变，且只有两个部件字节不同 ──
    zf2 = zipfile.ZipFile(io.BytesIO(patched_bytes))
    names_before = sorted(i.filename for i in zf.infolist())
    names_after = sorted(i.filename for i in zf2.infolist())
    if names_before != names_after:
        raise SystemExit("zip 部件集合变了 —— 拒绝写盘")
    differing = sorted(
        n for n in names_before if zf.read(n) != zf2.read(n)
    )
    report["differing_parts"] = differing
    if differing != sorted([part, "xl/styles.xml"]):
        raise SystemExit(
            f"改动的部件不止预期两个：{differing} —— 拒绝写盘"
        )

    # ── 前置校验③：单元格值一个都不许变 ──
    import openpyxl

    def _values(data: bytes) -> dict[str, Any]:
        wb = openpyxl.load_workbook(io.BytesIO(data))
        out_: dict[str, Any] = {}
        for ws in wb.worksheets:
            for row in ws.iter_rows():
                for cell in row:
                    if cell.value is not None:
                        out_[f"{ws.title}!{cell.coordinate}"] = cell.value
        wb.close()
        return out_

    if _values(original) != _values(patched_bytes):
        raise SystemExit("单元格值发生变化 —— 拒绝写盘")

    report["status"] = "would_patch" if not apply else "patched"
    print(f"[✓] 目标行 {rows}")
    print(f"[✓] style 重映射 {remap}（新增 {len(new_xfs)} 个 xf）")
    print(f"[✓] 改了 {changed} 个格的数字格式："
          f"{report['new_format']!r}（参照 {REFERENCE_COLUMN} 列）")
    print(f"[✓] 变动部件 = {differing}；全簿其余 number_format 与全部单元格值逐项不变")
    print(f"[✓] sha256 {report['sha256_before'][:16]}… → {report['sha256_after'][:16]}…")

    if apply:
        backup = template_path.with_suffix(template_path.suffix + ".bak")
        shutil.copy2(template_path, backup)
        template_path.write_bytes(patched_bytes)
        report["backup"] = str(backup.relative_to(_BACKEND))
        print(f"[apply] 已写盘，备份 → {report['backup']}")
        print("[apply] 🔴 接下来必须按顺序重生成：")
        print("        1) python backend/scripts/gen/generate_phase5_d1_contract.py --apply")
        print("        2) python backend/scripts/fix/fix_d1_republish_representation.py --apply")
    else:
        print("[check] 未写盘（--apply 才写）")
    return report


def main() -> int:
    parser = argparse.ArgumentParser(description=__doc__)
    mode = parser.add_mutually_exclusive_group(required=True)
    mode.add_argument("--check", action="store_true")
    mode.add_argument("--apply", action="store_true")
    parser.add_argument("--json", dest="json_path", default=None)
    args = parser.parse_args()
    report = run(apply=bool(args.apply))
    if args.json_path:
        Path(args.json_path).write_text(
            json.dumps(report, ensure_ascii=False, indent=2, default=str),
            encoding="utf-8",
        )
    return 0


if __name__ == "__main__":
    raise SystemExit(main())

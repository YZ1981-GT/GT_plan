# -*- coding: utf-8 -*-
"""净化 D3 预收账款权威模板的孤儿外链（过 OOXML external_relationships 门）。

G5-1 Phase 5 · 对齐 DEC-11（B60/D7 净化范式）：`backend/wp_templates/` 运行时只读，模板真
要升级必须显式净化 + 留 .bak 作门负例 + 重算 sentinel 重发布。

Task1 逐 zip 部件核实（tmp_d3_ext_probe）：
  - 3 个 externalLink 全是**断链孤儿**（Target 指向早已不存在的历史机器绝对路径：
    `Worksheet in 5440 Inventory` / `8240 COS breakdown` / `.../桌面/.../国贸_2005_随便.xls`），
    全 TargetMode=External。
  - **任何 worksheet 都没有 [n] 公式引用**（比 D7 更干净，无需中和公式）。
  - 受管 sheet 预收账款明细表D3-2 = sheet6.xml，零外链，124 格 / 22 merge。
  - defined name 44：#REF! 孤儿 27（保留，最小爆破面）+ bracket[n] 5（删外链后 dangle，删）。
  - workbook.xml 有 <externalReferences> 块（140 字符，删）。

净化 4 处（zip 级精准删除，比 D7 少一步「公式中和」，因 D3 无 [n] 公式格）：
  ① 删 3 个 xl/externalLinks/externalLink*.xml + 3 个 _rels/externalLink*.xml.rels
  ② [Content_Types].xml 删 3 个 <Override .../externalLink*.xml>
  ③ xl/_rels/workbook.xml.rels 删 3 个 externalLink <Relationship>
  ④ xl/workbook.xml 删 <externalReferences> 块 + 删含 [n] 的 defined name（#REF! 与其余保留）

用法（仓库根，PowerShell）：
  校验预演: & .venv/Scripts/python.exe backend/scripts/fix/sanitize_d3_template_external_links.py --check
  真净化:   & .venv/Scripts/python.exe backend/scripts/fix/sanitize_d3_template_external_links.py --apply

🔴 判成败一律用 openpyxl 打开 + D3-2 逐格 diff + 门判据，不看退出码。
"""
from __future__ import annotations

import hashlib
import io
import re
import shutil
import sys
import zipfile
from pathlib import Path

_REPO = Path(__file__).resolve().parents[3]
TEMPLATE = _REPO / "backend" / "wp_templates" / "D" / "D3 预收账款.xlsx"

_EXT_LINK_PART_RE = re.compile(r"^xl/externalLinks/externalLink\d+\.xml$")
_EXT_LINK_RELS_RE = re.compile(r"^xl/externalLinks/_rels/externalLink\d+\.xml\.rels$")
_WORKSHEET_RE = re.compile(r"^xl/worksheets/sheet\d+\.xml$")
_MANAGED_SHEET = "预收账款明细表D3-2"


def _strip_external_from_content_types(xml: str) -> str:
    return re.sub(
        r'<Override PartName="/xl/externalLinks/externalLink\d+\.xml"[^>]*/>',
        "",
        xml,
    )


def _strip_external_from_workbook_rels(xml: str) -> str:
    return re.sub(
        r'<Relationship [^>]*Target="externalLinks/externalLink\d+\.xml"[^>]*/>',
        "",
        xml,
    )


def _strip_external_from_workbook(xml: str) -> str:
    xml = re.sub(r"<externalReferences>.*?</externalReferences>", "", xml, flags=re.S)

    def _drop_bracket_dn(m: re.Match) -> str:
        return "" if re.search(r"\[\d+\]", m.group(0)) else m.group(0)

    xml = re.sub(r"<definedName [^>]*>.*?</definedName>", _drop_bracket_dn, xml, flags=re.S)
    return xml


def _neutralize_external_formulas(xml: str) -> str:
    """去掉含 [n] 的 <f>...</f>，保留同 <c> 内 <v> 缓存值。D3 实测 0 命中（无副作用）。"""
    return re.sub(r"<f>[^<]*\[\d+\][^<]*</f>", "", xml)


def sanitize_bytes(src: bytes) -> bytes:
    zin = zipfile.ZipFile(io.BytesIO(src))
    out = io.BytesIO()
    zout = zipfile.ZipFile(out, "w", zipfile.ZIP_DEFLATED)
    for info in zin.infolist():
        name = info.filename
        if _EXT_LINK_PART_RE.match(name) or _EXT_LINK_RELS_RE.match(name):
            continue  # ① 丢弃外链部件 + 其 rels
        data = zin.read(name)
        if name == "[Content_Types].xml":
            data = _strip_external_from_content_types(data.decode("utf-8")).encode("utf-8")
        elif name == "xl/_rels/workbook.xml.rels":
            data = _strip_external_from_workbook_rels(data.decode("utf-8")).encode("utf-8")
        elif name == "xl/workbook.xml":
            data = _strip_external_from_workbook(data.decode("utf-8")).encode("utf-8")
        elif _WORKSHEET_RE.match(name):
            # 兜底：D3 无 [n] 公式格，此处 0 命中；保留通用性以复用到 D5/D6
            data = _neutralize_external_formulas(data.decode("utf-8")).encode("utf-8")
        zout.writestr(name, data)
    zout.close()
    zin.close()
    return out.getvalue()


def _snapshot_managed(path_or_bytes) -> tuple[dict, list]:
    import openpyxl

    src = path_or_bytes if isinstance(path_or_bytes, (str, Path)) else io.BytesIO(path_or_bytes)
    wb = openpyxl.load_workbook(src, data_only=False)
    ws = wb[_MANAGED_SHEET]
    snap = {}
    for r in range(1, 40):
        for c in range(1, 28):
            v = ws.cell(r, c).value
            if v is not None:
                snap[f"{openpyxl.utils.get_column_letter(c)}{r}"] = str(v)
    merged = sorted(str(m) for m in ws.merged_cells.ranges)
    wb.close()
    return snap, merged


def _gate_ext_count(data: bytes) -> int:
    z = zipfile.ZipFile(io.BytesIO(data))
    n = 0
    for name in z.namelist():
        if name.endswith(".rels"):
            if 'TargetMode="External"' in z.read(name).decode("utf-8", "replace"):
                n += 1
    z.close()
    return n


def main() -> int:
    apply = "--apply" in sys.argv[1:]
    src = TEMPLATE.read_bytes()
    before_snap, before_merged = _snapshot_managed(TEMPLATE)
    before_sha = hashlib.sha256(src).hexdigest()

    out = sanitize_bytes(src)
    after_sha = hashlib.sha256(out).hexdigest()

    ext_before = _gate_ext_count(src)
    ext_after = _gate_ext_count(out)
    after_snap, after_merged = _snapshot_managed(out)
    diffs = [
        (k, before_snap.get(k), after_snap.get(k))
        for k in sorted(set(before_snap) | set(after_snap))
        if before_snap.get(k) != after_snap.get(k)
    ]

    print(f"before sha256={before_sha}")
    print(f"after  sha256={after_sha}")
    print(f"external .rels: before={ext_before} after={ext_after} → 门:{'PASS' if ext_after == 0 else 'REJECT'}")
    print(f"D3-2 managed cells before={len(before_snap)} after={len(after_snap)} diffs={len(diffs)}")
    for k, b, a in diffs[:30]:
        print(f"  {k}: {b!r} -> {a!r}")
    print(f"D3-2 merged same={before_merged == after_merged}")

    ok = ext_after == 0 and len(diffs) == 0 and before_merged == after_merged
    if not ok:
        print("[FAIL] 净化不满足判据（门未过 / 受管 sheet 有 diff / merge 变了）—— 不写盘")
        return 1

    if apply:
        bak = TEMPLATE.with_suffix(TEMPLATE.suffix + ".preclean.bak")
        if not bak.exists():
            shutil.copy2(TEMPLATE, bak)
            print(f"[apply] 备份门负例 → {bak.name}")
        TEMPLATE.write_bytes(out)
        print(f"[apply] 已净化写盘 {TEMPLATE.name}  new_sha256={after_sha}")
    else:
        print("[check] 判据全过（门 PASS + D3-2 0 diff + merge 不变）；--apply 才写盘")
    return 0


if __name__ == "__main__":
    raise SystemExit(main())

# -*- coding: utf-8 -*-
"""净化 M2/M3/M4/M7/M10 权威模板的外部关系（过首版发布 OOXML `external_relationships` 门）。

spec: `m2-m3-m4-m7-m10-bidirectional-pipeline`

逐册现查（2026-10-07，zip 内逐部件扫描）：

| 册 | 外链部件 | `[n]` 公式 | `[n]` definedName | 外部 hyperlink |
|---|---|---|---|---|
| M2 实收资本 | 3（×2=6 文件） | 0 | 9 | 0 |
| M3 库存股 | 3（×2=6 文件） | 0 | 9 | 0 |
| M4 资本公积 | 3（×2=6 文件） | 0 | 9 | 0 |
| M7 专项储备 | 3（×2=6 文件） | 0 | 9 | 0 |
| M10 其他权益工具 | 3（×2=6 文件） | 0 | 13 | 1 |

处置（照 I/J/D 系同口径）：
  ① 删 `xl/externalLinks/*` 部件与其 `_rels`
  ② `[Content_Types].xml` 删对应 Override
  ③ `xl/_rels/workbook.xml.rels` 删对应 Relationship
  ④ `xl/workbook.xml` 删 `<externalReferences>` 段和含 `[n]` 的 definedName
  ⑤ sheet rels 里 `TargetMode="External"` 的 Relationship 删除
  ⑥ sheet xml 里对应的 `<hyperlink r:id=…/>` 删除

用法：
  & .venv/Scripts/python.exe backend/scripts/fix/sanitize_m_cycle_template_external_links.py          # 预演
  & .venv/Scripts/python.exe backend/scripts/fix/sanitize_m_cycle_template_external_links.py --apply  # 写盘 + .bak
"""
from __future__ import annotations

import hashlib
import io
import re
import sys
import zipfile
from dataclasses import dataclass
from pathlib import Path

_REPO = Path(__file__).resolve().parents[3]
sys.path.insert(0, str(_REPO / "backend"))
_TPL_DIR = _REPO / "backend" / "wp_templates" / "M"


@dataclass(frozen=True)
class _Target:
    file_name: str
    managed_sheet: str
    expect_ext_parts: int        # externalLink xml + rels 文件数
    expect_formulas: int         # [n] 公式
    expect_defined_names: int    # [n] definedName
    expect_hyperlinks: int       # TargetMode="External" hyperlink


TARGETS: tuple[_Target, ...] = (
    _Target("M2 实收资本（股本）.xlsx", "明细表（非上市公司）M2-2", 6, 0, 9, 0),
    _Target("M3 库存股.xlsx", "明细表M3-2", 6, 0, 9, 0),
    _Target("M4 资本公积.xlsx", "明细表M4-2", 6, 0, 9, 0),
    _Target("M7 专项储备.xlsx", "明细表M7-2", 6, 0, 9, 0),
    _Target("M10 其他权益工具.xlsx", "明细表M10-2", 6, 0, 13, 1),
)

_EXT_LINK_PART_RE = re.compile(r"^xl/externalLinks/externalLink\d+\.xml$")
_EXT_LINK_RELS_RE = re.compile(r"^xl/externalLinks/_rels/externalLink\d+\.xml\.rels$")
_EXT_FORMULA_RE = re.compile(r"<f>[^<]*\[\d+\][^<]*</f>")
_EXT_REL_RE = re.compile(r'<Relationship\b[^>]*TargetMode="External"[^>]*/>')


def _strip_workbook(xml: str) -> tuple[str, int]:
    """删 <externalReferences> 段和含 [n] 的 definedName。"""
    xml = re.sub(r"<externalReferences>.*?</externalReferences>", "", xml, flags=re.S)
    dropped = 0

    def _drop(m: re.Match) -> str:
        nonlocal dropped
        if re.search(r"\[\d+\]", m.group(0)):
            dropped += 1
            return ""
        return m.group(0)

    xml = re.sub(r"<definedName\b[^>]*>.*?</definedName>", _drop, xml, flags=re.S)
    xml = re.sub(r"<definedNames>\s*</definedNames>", "", xml)
    return xml, dropped


def sanitize_bytes(src: bytes) -> tuple[bytes, dict[str, int]]:
    zin = zipfile.ZipFile(io.BytesIO(src))
    stats = {"dropped_parts": 0, "neutralized_formulas": 0,
             "dropped_defined_names": 0, "dropped_hyperlinks": 0}

    # 先找每个 sheet 被删的外部 hyperlink rId
    drop_rids: dict[str, set[str]] = {}
    for name in zin.namelist():
        m = re.match(r"^xl/worksheets/_rels/(sheet\d+\.xml)\.rels$", name)
        if not m:
            continue
        rels = zin.read(name).decode("utf-8")
        ids = set()
        for rel_match in _EXT_REL_RE.finditer(rels):
            id_match = re.search(r'Id="([^"]+)"', rel_match.group(0))
            if id_match:
                ids.add(id_match.group(1))
        if ids:
            drop_rids[f"xl/worksheets/{m.group(1)}"] = ids

    out = io.BytesIO()
    zout = zipfile.ZipFile(out, "w", zipfile.ZIP_DEFLATED)
    for info in zin.infolist():
        name = info.filename
        # 删 externalLinks 部件
        if _EXT_LINK_PART_RE.match(name) or _EXT_LINK_RELS_RE.match(name):
            stats["dropped_parts"] += 1
            continue
        data = zin.read(name)
        if name == "[Content_Types].xml":
            text = re.sub(
                r'<Override PartName="/xl/externalLinks/externalLink\d+\.xml"[^>]*/>', "",
                data.decode("utf-8"),
            )
            data = text.encode("utf-8")
        elif name == "xl/_rels/workbook.xml.rels":
            data = re.sub(
                r'<Relationship [^>]*Target="externalLinks/externalLink\d+\.xml"[^>]*/>', "",
                data.decode("utf-8"),
            ).encode("utf-8")
        elif name == "xl/workbook.xml":
            text, n = _strip_workbook(data.decode("utf-8"))
            stats["dropped_defined_names"] += n
            data = text.encode("utf-8")
        elif re.match(r"^xl/worksheets/_rels/sheet\d+\.xml\.rels$", name):
            text = data.decode("utf-8")
            text = _EXT_REL_RE.sub("", text)
            data = text.encode("utf-8")
        elif name.startswith("xl/worksheets/sheet") and name.endswith(".xml"):
            text = data.decode("utf-8")
            hits = len(_EXT_FORMULA_RE.findall(text))
            if hits:
                stats["neutralized_formulas"] += hits
                text = _EXT_FORMULA_RE.sub("", text)
            for rid in sorted(drop_rids.get(name, ())):
                text, n = re.subn(
                    r'<hyperlink\b[^>]*r:id="' + re.escape(rid) + r'"[^>]*/>', "", text
                )
                stats["dropped_hyperlinks"] += n
            text = re.sub(r"<hyperlinks>\s*</hyperlinks>", "", text)
            data = text.encode("utf-8")
        zout.writestr(info, data)
    zout.close()
    zin.close()
    return out.getvalue(), stats


def _snapshot(data: bytes, sheet: str) -> dict[str, str]:
    import openpyxl
    wb = openpyxl.load_workbook(io.BytesIO(data), data_only=False)
    ws = wb[sheet]
    snap = {}
    for row in ws.iter_rows(min_row=1, max_row=ws.max_row, max_col=ws.max_column):
        for c in row:
            if c.value is not None:
                text = getattr(c.value, "text", None)
                val = f"{{array {getattr(c.value, 'ref', '')}}}{text}" if text is not None else str(c.value)
                snap[c.coordinate] = val
    wb.close()
    return snap


def _gate(data: bytes) -> tuple[bool, str]:
    """真跑生产 OOXML 安全门。"""
    import tempfile
    from app.services.workpaper_sync.ooxml_security import validate_ooxml_artifact
    with tempfile.TemporaryDirectory() as tmp:
        p = Path(tmp) / "probe.xlsx"
        p.write_bytes(data)
        try:
            validate_ooxml_artifact(p, document_type="xlsx")
            return True, "PASS"
        except Exception as exc:
            return False, f"REJECT {getattr(exc, 'gate', type(exc).__name__)}"


def run_one(t: _Target, *, apply: bool) -> bool:
    path = _TPL_DIR / t.file_name
    bak = path.with_suffix(path.suffix + ".preclean.bak")
    src = bak.read_bytes() if bak.exists() else path.read_bytes()
    out, stats = sanitize_bytes(src)

    # 受管 sheet 逐格对比
    b_snap = _snapshot(src, t.managed_sheet)
    a_snap = _snapshot(out, t.managed_sheet)
    diffs = [k for k in sorted(set(b_snap) | set(a_snap)) if b_snap.get(k) != a_snap.get(k)]

    gate_before = _gate(src)
    gate_after = _gate(out)

    expected = {
        "dropped_parts": t.expect_ext_parts,
        "neutralized_formulas": t.expect_formulas,
        "dropped_defined_names": t.expect_defined_names,
        "dropped_hyperlinks": t.expect_hyperlinks,
    }
    ok = True
    sha_before = hashlib.sha256(src).hexdigest()[:16]
    sha_after = hashlib.sha256(out).hexdigest()[:16]

    print(f"\n{'='*60}")
    print(f"  {t.file_name}")
    print(f"  SHA256: {sha_before}... -> {sha_after}...")
    print(f"  Gate: {gate_before[1]} -> {gate_after[1]}")
    print(f"  Stats: {stats}")
    print(f"  Expected: {expected}")
    print(f"  Managed sheet diffs: {len(diffs)}")
    if diffs:
        for d in diffs[:5]:
            print(f"    {d}: {b_snap.get(d, '<missing>')} -> {a_snap.get(d, '<missing>')}")

    if stats != expected:
        print(f"  ❌ Stats mismatch!")
        ok = False
    if diffs:
        print(f"  ❌ {len(diffs)} cells differ on managed sheet!")
        ok = False
    if not gate_before[1].startswith("REJECT"):
        print(f"  ❌ Original should be REJECTED!")
        ok = False
    if gate_after[1] != "PASS":
        print(f"  ❌ Sanitized should PASS!")
        ok = False

    if ok and apply:
        if not bak.exists():
            bak.write_bytes(src)
            print(f"  备份: {bak.name}")
        path.write_bytes(out)
        new_sha = hashlib.sha256(out).hexdigest()
        print(f"  ✅ 已写入: {path.name} SHA256={new_sha}")
    elif ok:
        print(f"  ✅ 预演通过（加 --apply 写盘）")
    else:
        print(f"  ❌ 校验未通过，不写盘")

    return ok


def main():
    import argparse
    parser = argparse.ArgumentParser(description=__doc__ or "")
    parser.add_argument("--apply", action="store_true", help="写盘 + .bak")
    args = parser.parse_args()

    results = []
    for t in TARGETS:
        results.append(run_one(t, apply=args.apply))

    passed = sum(results)
    print(f"\n{'='*60}")
    print(f"  结果: {passed}/{len(results)} 通过")
    if not all(results):
        sys.exit(1)


if __name__ == "__main__":
    main()

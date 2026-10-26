# -*- coding: utf-8 -*-
r"""按正确口径修权威模板里的公式错误（spec oo-writeback-template-parity-closure R1）。

用户 2026-10-26：「模板本身算错的几处，这个你看下实际情况，如果真的模板错误就按正确的来；
D5-2 改成 L+M+N（包含重分类），可能也是模板错了吧」。逐格核实依据见 spec design §2。

做法沿用 ``fix_d3_template_footer_defects.py``：直接改 worksheet XML（独立公式改文本、
共享公式只改主格、空格补 ``<f>``、内联文本改字），其余 part 逐字节不动；判成败一律
openpyxl 逐格比对：**只有清单内的格变化，且变成期望值**，merged 区不变。

用法（backend 目录）：
  校验预演: ..\.venv\Scripts\python.exe scripts/fix/fix_template_formula_defects.py --check
  真修复:   ..\.venv\Scripts\python.exe scripts/fix/fix_template_formula_defects.py --apply
已修过的格视为完成（幂等）：现状 == 期望值的格跳过。
"""
from __future__ import annotations

import hashlib
import io
import re
import sys
import warnings
import zipfile
from dataclasses import dataclass
from pathlib import Path

from openpyxl.utils import column_index_from_string, get_column_letter

ROOT = Path(__file__).resolve().parents[2]  # backend/
TEMPLATES = ROOT / "wp_templates"


@dataclass(frozen=True)
class Fix:
    template: str
    sheet: str
    coord: str
    old: str | None  # openpyxl 读到的现状（公式带 =；None = 空格）
    new: str | None  # 期望值（公式带 =；否则是文本；None = 清成空的录入格，保留样式）


def _col(column: str, rows: range, old: str | None, new: str, template: str, sheet: str) -> list[Fix]:
    out = []
    for r in rows:
        out.append(Fix(template, sheet, f"{column}{r}",
                       None if old is None else old.format(r=r), new.format(r=r)))
    return out


D5, F1, K3 = "D/D5 应收款项融资.xlsx", "F/F1 预付账款.xlsx", "K/K3 其他应付款.xlsx"
N1, N3, N4 = "N/N1 递延所得税资产.xlsx", "N/N3 递延所得税负债.xlsx", "N/N4 税金及附加.xlsx"
H2, G8, D6 = "H/H2 在建工程.xlsx", "G/G8 其他权益工具投资.xlsx", "D/D6 合同资产.xlsx"
K1, D1, G5 = "K/K1 其他应收款.xlsx", "D/D1 应收票据.xlsx", "G/G5 长期应收款.xlsx"
I3, H3 = "I/I3 商誉.xlsx", "H/H3 投资性房地产.xlsx"
G7, K6, L4 = "G/G7 长期股权投资.xlsx", "K/K6 持有待售资产和负债.xlsx", "L/L4 应付债券.xlsx"
D2I, D7, G10 = "D/D2-6至D2-13  应收账款 -检查（Leap应对措施-检查）.xlsx", "D/D7 合同负债.xlsx", "G/G10 交易性金融负债.xlsx"
D2A, G13 = "D/D2-1至D2-4  应收账款- 审定表明细表（Leap-常规程序）.xlsx", "G/G13 公允价值变动收益.xlsx"
D4_BOOKS = ("D/D4 收入底稿.xlsx", "D/D4-22至D4-32营业收入-IPO 上市 新三板 重组 舞弊应对.xlsx")
N12S, N32S = "递延所得税资产明细表N1-2", "递延所得税负债明细表N3-2"


def fixes() -> list[Fix]:
    out: list[Fix] = []
    # §2-1 D5-2 期末审定 = 期末未审 + 账项 + 重分类（同结构六张表都是这样）
    out += _col("O", range(12, 17), "=J{r}+N{r}+M{r}", "=L{r}+M{r}+N{r}", D5, "应收款项融资明细表D5-2")
    # §2-2/3 F1-2、K3-2 同型：审定数从期末余额起算漏了被审计单位重分类
    out += _col("X", range(14, 35), "=O{r}+V{r}+W{r}", "=Q{r}+V{r}+W{r}", F1, "明细表F1-2")
    out += _col("T", range(10, 31), "=O{r}+R{r}+S{r}", "=Q{r}+R{r}+S{r}", K3, "明细表K3-2")
    # §2-4 N1-2：审定数漏重分类调整；G / M 列无合计
    out += _col("H", range(11, 41), "=F{r}+E{r}", "=E{r}+F{r}+G{r}", N1, N12S)
    out += _col("N", range(11, 41), "=L{r}+K{r}", "=K{r}+L{r}+M{r}", N1, N12S)
    out += [Fix(N1, N12S, "G42", None, "=SUM(G11:G41)"), Fix(N1, N12S, "M42", None, "=SUM(M11:M41)")]
    # §2-5 N3-2：期初审定 F+O3…（O3 起是空格）、期末审定漏重分类、E 列合计区间整块错、表头错字
    out += [Fix(N3, N32S, f"H{r}", f"=F{r}+O{r - 8}", f"=E{r}+F{r}+G{r}") for r in range(11, 22)]
    out += _col("N", range(11, 22), "=L{r}+K{r}", "=K{r}+L{r}+M{r}", N3, N32S)
    out += [Fix(N3, N32S, "E23", "=SUM(E3:O22)", "=SUM(E11:E22)"),
            Fix(N3, N32S, "G23", None, "=SUM(G11:G22)"), Fix(N3, N32S, "M23", None, "=SUM(M11:M22)"),
            Fix(N3, N32S, "H10", "递延所得税资产期初审定数", "递延所得税负债期初审定数")]
    # §2-6 N4-2：首行第三项引错格；D 列合计区间整块错
    out += [Fix(N4, "税金及附加明细表N4-2", "E9", "=B9+C9+N4", "=B9+C9+D9"),
            Fix(N4, "税金及附加明细表N4-2", "D19", "=SUM(D4:N18)", "=SUM(D9:D18)")]
    # §2-7 H2-2：期末三处都错成右边一列（「其中：累计资本化」）；比例除数是预算 B
    out += _col("D", range(13, 21), "=IF(AH{r}=0,0,AH{r}/B{r})", "=IF(B{r}=0,0,AG{r}/B{r})", H2, "明细表H2-2")
    out += _col("AV", range(13, 21), "=R{r}-AL{r}", "=Q{r}-AL{r}", H2, "明细表H2-2")
    out += _col("AW", range(13, 21), "=AH{r}-AS{r}", "=AG{r}-AS{r}", H2, "明细表H2-2")
    # §2-8 G8-2：M11 多加了 L（M = Q − E，L 是权益内部结转）、P 漏加 K、R/T 缺公式
    g8 = "明细表G8-2"
    out += [Fix(G8, g8, "M11", "=SUM(I11:L11)", "=SUM(I11:K11)"), Fix(G8, g8, "P11", "=D11+J11", "=D11+J11+K11")]
    out += _col("P", range(13, 21), "=D{r}+J{r}", "=D{r}+J{r}+K{r}", G8, g8)
    out += _col("R", range(13, 21), None, "=F{r}+J{r}+L{r}", G8, g8)
    # 合计行 K / L / R 三列没有 SUM（同行其余数值列都有）
    out += [Fix(G8, g8, f"{c}21", None, f"=SUM({c}11:{c}20)") for c in ("K", "L", "R")]
    out += [Fix(G8, g8, "T12", None, "=Q12+S12")]
    # §2-9 D5-4：G / I / J 只有首行有公式
    d54 = "应收款项融资公允价值测算表D5-4"
    out += _col("G", range(13, 17), None, "=F{r}-E{r}", D5, d54)
    out += _col("I", range(13, 17), None, "=D{r}*H{r}*G{r}/365", D5, d54)
    out += _col("J", range(13, 17), None, "=D{r}-I{r}", D5, d54)
    # §2-10/11 小计行乱写
    out += [Fix(D6, "合同资产减值准备明细表D6-3", f"N{r}", f"=E{r}+SUM(H{r}:I{r})-SUM(J{r}:L{r})",
                f"=K{r}+L{r}+M{r}") for r in (12, 17)]
    out += [Fix(K1, "坏账准备明细表K1-3", "L11", "=D11+SUM(G11:G11)-SUM(I11:K11)", "=J11+K11")]
    # §2-12 D1-7：同列其他行取 C 列日期
    out += [Fix(D1, "应收票据备查簿核对D1-7", "Q23", "=IF(YEAR(A23)=YEAR($B$10)-1,H23,0)",
                "=IF(YEAR(C23)=YEAR($B$10)-1,H23,0)")]
    # §2-13 D4-27：M17:M19 缺「总计」（两本 D4 工作簿同一张表）
    for book in D4_BOOKS:
        out += _col("M", range(17, 20), None, "=SUM(C{r}:L{r})", book, "识别未披露的关联方D4-27")
    # §2-14 已登记未修：I3-2 引错区间 + 首行漏公式、H3-2 公允 R 列缺公式、G5-2 段合计漏小计、G5-1 越界
    i32 = "明细表I3-2"
    out += [Fix(I3, i32, f"{c}23", f"=SUM({c}27:{c}30)", f"=SUM({c}14:{c}22)") for c in ("AA", "AB", "AC", "AD")]
    out += [Fix(I3, i32, "I14", None, "=SUM(D14:E14)-G14")]
    out += _col("R", range(13, 27), None, "=P{r}+Q{r}", H3, "明细表（公允价值模式）H3-2")
    g52 = "余额明细表G5-2"
    for total, (a, b, c, d) in ((40, (18, 25, 32, 39)), (72, (50, 57, 64, 71)), (104, (82, 89, 96, 103))):
        for col in (get_column_letter(i) for i in range(column_index_from_string("D"), column_index_from_string("R") + 1)):
            out.append(Fix(G5, g52, f"{col}{total}", f"={col}{a}+{col}{b}+{col}{d}",
                           f"={col}{a}+{col}{b}+{col}{c}+{col}{d}"))
    out += [Fix(G5, "审定表G5-1", "B35", "=B9-B225", "=B9-B22")]
    # §2-15 同行跨表单格引用指错行（scan_row_xref 普查后逐条核实）
    # G5-1 期末账项 / 重分类取了明细表首个数据行 R13，同行 B..F 都取小计行 R18
    out += [Fix(G5, "审定表G5-1", "G9", "='余额明细表G5-2'!K13", "='余额明细表G5-2'!K18"),
            Fix(G5, "审定表G5-1", "H9", "='余额明细表G5-2'!L13", "='余额明细表G5-2'!L18")]
    # G7-1 减值准备「重分类调整 · 对子公司」：期初取了数据行 M62，增减抄成账项调整的 N/O 列
    # （同块 R40/R41 是 M/P/Q 小计；明细表 R79 S=H+N+P、T=I+O+Q ⇒ P/Q 是重分类增减）
    g71, g72 = "长期股权投资审定表G7-1", "'明细表G7-2'"
    out += [Fix(G7, g71, "C39", f"={g72}!M62", f"={g72}!M63"),
            Fix(G7, g71, "D39", f"={g72}!N63", f"={g72}!P63"),
            Fix(G7, g71, "E39", f"={g72}!O63", f"={g72}!Q63")]
    # K6 附注（上市）上年年末数整列错位一行：E9~E11 指向下一行，E8 应与 B8 一样是小计
    k6n, k62 = "附注披露信息（上市公司）", "'明细表K6-2'"
    out += [Fix(K6, k6n, "E8", f"={k62}!L12", "=SUM(E9:E11)")]
    out += [Fix(K6, k6n, f"E{r}", f"={k62}!L{r + 4}", f"={k62}!L{r + 3}") for r in (9, 10, 11)]
    # L4 附注（上市）「是否违约」：可转换债券两行取到了小计行 / 区标题行（同行其余列取 R19/R20）
    l4n, l42 = "附注披露信息核对（上市公司）", "'应付债券明细表L4-2'"
    out += [Fix(L4, l4n, "H30", f"={l42}!I17", f"={l42}!I19"),
            Fix(L4, l4n, "H31", f"={l42}!I18", f"={l42}!I20")]
    # §2-16 L4-3 合计行期初调整四列抄成了「期初 + 增加 − 减少」的错位交叉式（已登记未修）：
    # 数据行 T/U/V/W 都是录入格，合计应与同行 L..Q、X..AK 一样逐列 SUM
    l43 = "划分为金融负债的其他金融工具明细表L4-3"
    for col, a, b, c in (("T", "N", "P", "R"), ("U", "O", "Q", "S"), ("V", "P", "R", "T"), ("W", "Q", "S", "U")):
        out.append(Fix(L4, l43, f"{col}18", f"={a}18+{b}18-{c}18", f"=SUM({col}13:{col}17)"))
    # ── 第二批（按模板对齐前端时逐表核出，design §2 #20~#25）──
    # §2-20 D1-1 银行承兑汇票期初三列取了 D1-2 合计行 R14（同行期末 F..H 取 R11、下一行取 R13；
    # 取合计会让小计 B10 把商业承兑汇票算两遍）
    d11, d12 = "审定表D1-1", "'原值明细表（按类别）D1-2'"
    out += [Fix(D1, d11, f"{c}8", f"={d12}!{c}14", f"={d12}!{c}11") for c in ("B", "C", "D")]
    # §2-21 D1-15 ④「期末坏账准备账面余额」是被审计单位已记录的坏账准备（录入；A8「与被审计单位已记录的
    # 坏账准备进行比较」），模板写成 ①−③（账面价值）⇒ ⑤=③−④ 恒为 2×③−①；同结构 D2-9 的 ④ 是录入格
    d115 = "应收票据坏账准备测试表D1-15"
    out += [Fix(D1, d115, f"E{r}", f"=B{r}-D{r}", None) for r in (*range(14, 18), *range(22, 25))]
    out += [Fix(D1, d115, f"E{r}", "④=①-③", "④") for r in (13, 21)]
    # §2-22 D2-9 单项区差异写成 ①−④（表头 ⑤=③−④，两个组合区各行都是 D−E）；
    # 小计行「预期信用损失率」列在加总比例、「应计提」列在算比例，两格互换
    d29 = "应收坏账准备测算D2-9"
    out += _col("F", range(14, 19), "=B{r}-E{r}", "=D{r}-E{r}", D2I, d29)
    out += [Fix(D2I, d29, "C19", "=SUM(C14:C18)", '=IF(B19=0,"",D19/B19)'),
            Fix(D2I, d29, "D19", '=IF(B19=0,"",C19/B19)', "=SUM(D14:D18)")]
    # §2-23 D7-6 合同负债是贷方科目：期末 = 期初 − 借方 + 贷方（同册 D7-2 P=F−N+O）；模板按资产方向写成 C+D−E
    out += _col("F", range(12, 15), "=C{r}+D{r}-E{r}", "=C{r}-D{r}+E{r}", D7, "关联方关系及交易检查表D7-6")
    # §2-24 D7-1 期初重分类小计多含表头行 R7（同行其余列 SUM(x8:x13)；SUM 忽略文字，数值不变）
    out += [Fix(D7, "审定表D7-1", "D14", "=SUM(D7:D13)", "=SUM(D8:D13)")]
    # §2-25 G10-1「发行的交易性债券」：变动额 / 变动率取了上一行小计 G8，账面余额 G29 也取了 G8
    # （同行 D..F、H 与其他行都取本行 R9）
    g101 = "审定表G10-1"
    out += [Fix(G10, g101, "H9", "=G8-D9", "=G9-D9"),
            Fix(G10, g101, "I9", '=IF(AND(D9=0,G8=0),"",IF(AND(D9=0,G8>0),1,H9/D9))',
                '=IF(AND(D9=0,G9=0),"",IF(AND(D9=0,G9>0),1,H9/D9))'),
            Fix(G10, g101, "G29", "=G8+G19", "=G9+G19")]
    # §2-26 D2-3 组合区数据行是 R18~R21（四行都有 E/K/N 公式），小计 R17 只有 E/K/N 求和到 R21、其余 10 列
    # 漏了 R21；期末审定 N18:N21 漏加重分类 M（单项区 N13:N16 与同结构 D1-4 都是 K+L+M）
    d23 = "坏账准备明细表D2-3"
    out += [Fix(D2A, d23, f"{c}17", f"=SUM({c}18:{c}20)", f"=SUM({c}18:{c}21)") for c in "BCDFGHIJLM"]
    out += _col("N", range(18, 22), "=K{r}+L{r}", "=K{r}+L{r}+M{r}", D2A, d23)
    # §2-27 D2-10 方法 2 的前瞻性调整 C 列逐行取 G64~G69，末行 C90 取到了表外的 G70（R70 是说明文字行）
    out += [Fix(D2I, "预期信用损失的计量测试D2-10", "C90", "=G70", "=G69")]
    # §2-28 G13-1 合计行上期审定取了「衍生金融资产」G9（同行其余列取小计 R7）；R7 变动率分母也是 G9
    g131 = "审定表G13-1"
    out += [Fix(G13, g131, "G17", "=SUM(G9,G10,G13,G15:G16)", "=SUM(G7,G10,G13,G15:G16)"),
            Fix(G13, g131, "I7", '=IF(AND(G9=0,D7=0),"",IF(AND(G9=0,D7>0),1,H7/G9))',
                '=IF(AND(G7=0,D7=0),"",IF(AND(G7=0,D7>0),1,H7/G7))')]
    return out


# ── OOXML 层 ────────────────────────────────────────────────────────────────

_CELL = r'<c r="{coord}"(?P<attrs>(?:\s[^>/]*)?)(?:/>|>(?P<body>.*?)</c>)'
_F_TEXT = re.compile(r"<f(?P<fattrs>[^>]*)>(?P<text>.*?)</f>", re.S)
_F_EMPTY = re.compile(r"<f(?P<fattrs>[^>]*)/>")


def _sheet_parts(data: bytes) -> dict[str, str]:
    with zipfile.ZipFile(io.BytesIO(data)) as zf:
        wb = zf.read("xl/workbook.xml").decode("utf-8")
        rels = zf.read("xl/_rels/workbook.xml.rels").decode("utf-8")
    targets = {}
    for rel in re.findall(r"<Relationship\b[^>]*>", rels):
        rid, tgt = re.search(r'\bId="([^"]+)"', rel), re.search(r'\bTarget="([^"]+)"', rel)
        if rid and tgt:
            targets[rid.group(1)] = tgt.group(1).lstrip("/")
    out = {}
    for name, rid in re.findall(r'<sheet [^>]*name="([^"]+)"[^>]*r:id="([^"]+)"', wb):
        target = targets.get(rid, "")
        out[name.replace("&amp;", "&")] = target if target.startswith("xl/") else f"xl/{target}"
    return out


def _esc(text: str) -> str:
    return text.replace("&", "&amp;").replace("<", "&lt;").replace(">", "&gt;")


# 共享公式主格：开标签不能是自闭合（否则会从前面的空格一路吃到主格），且不跨格
_MASTER = r'<c r="(?P<coord>[A-Z]+\d+)"(?:\s[^>/]*)?>(?:(?!</c>|<c\b).)*?<f(?P<fattrs>[^>]*\bsi="{si}"[^>]*)>'


def _shared_master(xml: str, si: str) -> tuple[str, str] | None:
    """返回 (主格坐标, ref)；从属格的 <f> 没有 ref，跳过。"""
    for m in re.finditer(_MASTER.format(si=si), xml, re.S):
        ref = re.search(r'\bref="([^"]+)"', m.group("fattrs"))
        if ref:
            return m.group("coord"), ref.group(1)
    return None


_DEPENDENT = r'<c r="(?P<coord>[A-Z]+\d+)"(?:\s[^>/]*)?>(?:(?!</c>|<c\b).)*?<f[^>]*\bsi="{si}"[^>]*/>'


def _shared_members(xml: str, si: str, master: str) -> set[str]:
    """共享公式组里真正用这条公式的格（主格 + 从属格）。

    ``ref`` 只是外接矩形：区内自带公式的格（如 D2-3 的 E17）不属于该组，不能按 ``ref`` 判。
    """
    return {master} | {m.group("coord") for m in re.finditer(_DEPENDENT.format(si=si), xml, re.S)}


def _ref_cells(ref: str) -> set[str]:
    a, _, b = ref.partition(":")
    b = b or a
    ma, mb = re.match(r"([A-Z]+)(\d+)", a), re.match(r"([A-Z]+)(\d+)", b)
    cols = range(column_index_from_string(ma.group(1)), column_index_from_string(mb.group(1)) + 1)
    return {f"{get_column_letter(c)}{r}" for c in cols for r in range(int(ma.group(2)), int(mb.group(2)) + 1)}


def _assert_group_in_targets(xml: str, fix: Fix, fattrs: str, targets: dict[str, Fix]) -> None:
    """改共享公式主格 = 改整组：组里每个格都必须在清单内（否则会把清单外的格一起改掉）。"""
    ref = re.search(r'\bref="([^"]+)"', fattrs)
    si = re.search(r'\bsi="(\d+)"', fattrs)
    if ref is None or si is None:
        return
    outside = _shared_members(xml, si.group(1), fix.coord) - set(targets)
    if outside:
        raise RuntimeError(f"{fix.sheet}!{fix.coord} 是共享公式主格（{ref.group(1)}），"
                           f"组内有清单外的格 {sorted(outside)}")


def _patch_sheet(xml: str, items: list[Fix]) -> str:
    targets = {f.coord: f for f in items}
    for fix in items:
        m = re.search(_CELL.format(coord=fix.coord), xml, re.S)
        if m is None:
            raise RuntimeError(f"{fix.sheet}!{fix.coord} 在 XML 里不存在（不凭空造格）")
        attrs, body = m.group("attrs") or "", m.group("body") or ""
        if fix.new is None:  # 清成空的录入格：去掉公式与缓存值，只留样式
            f_text = _F_TEXT.search(body)
            fattrs = f_text.group("fattrs") if f_text else (_F_EMPTY.search(body).group("fattrs")
                                                          if _F_EMPTY.search(body) else "")
            _assert_group_in_targets(xml, fix, fattrs, targets)
            clean = re.sub(r'\s+t="[^"]*"', "", attrs)
            new_cell = f'<c r="{fix.coord}"{clean}/>'
        elif not fix.new.startswith("="):  # 文本格：内联字符串原地改；共享字符串格转成内联字符串
            if 't="inlineStr"' in attrs:
                new_cell = f'<c r="{fix.coord}"{attrs}><is><t>{_esc(fix.new)}</t></is></c>'
            elif 't="s"' in attrs:
                inline = attrs.replace('t="s"', 't="inlineStr"')
                new_cell = f'<c r="{fix.coord}"{inline}><is><t>{_esc(fix.new)}</t></is></c>'
            else:
                raise RuntimeError(f"{fix.sheet}!{fix.coord} 不是文本格，拒绝改写")
        else:
            f_text, f_empty = _F_TEXT.search(body), _F_EMPTY.search(body)
            if f_text is not None:
                fattrs = f_text.group("fattrs")
                _assert_group_in_targets(xml, fix, fattrs, targets)
                new_body = body.replace(f_text.group(0), f"<f{fattrs}>{_esc(fix.new[1:])}</f>", 1)
                new_cell = f'<c r="{fix.coord}"{attrs}>{new_body}</c>'
            elif f_empty is not None:
                si = re.search(r'\bsi="(\d+)"', f_empty.group("fattrs"))
                master = _shared_master(xml, si.group(1)) if si else None
                if master is None or master[0] not in targets or fix.coord not in _ref_cells(master[1]):
                    raise RuntimeError(f"{fix.sheet}!{fix.coord} 是共享公式从属格，主格不在清单内")
                continue  # 随主格一起改
            else:  # 空格 / 只有值：补公式
                clean = re.sub(r'\s+t="[^"]*"', "", attrs)
                new_cell = f'<c r="{fix.coord}"{clean}><f>{_esc(fix.new[1:])}</f></c>'
        xml = xml[: m.start()] + new_cell + xml[m.end():]
    return xml


def patch_bytes(data: bytes, items: list[Fix]) -> bytes:
    parts = _sheet_parts(data)
    by_part: dict[str, list[Fix]] = {}
    for fix in items:
        by_part.setdefault(parts[fix.sheet], []).append(fix)
    zin = zipfile.ZipFile(io.BytesIO(data))
    out = io.BytesIO()
    with zipfile.ZipFile(out, "w", zipfile.ZIP_DEFLATED) as zout:
        for info in zin.infolist():
            payload = zin.read(info.filename)
            if info.filename in by_part:
                payload = _patch_sheet(payload.decode("utf-8"), by_part[info.filename]).encode("utf-8")
            zout.writestr(info, payload)
    zin.close()
    return out.getvalue()


def _snapshot(data: bytes) -> dict[str, tuple[dict[str, str], list[str]]]:
    import openpyxl

    warnings.filterwarnings("ignore")
    wb = openpyxl.load_workbook(io.BytesIO(data), data_only=False)
    snap = {}
    for ws in wb.worksheets:
        cells = {c.coordinate: str(c.value) for row in ws.iter_rows() for c in row if c.value is not None}
        snap[ws.title] = (cells, sorted(str(m) for m in ws.merged_cells.ranges))
    wb.close()
    return snap


def run(apply: bool, templates: Path = TEMPLATES) -> int:
    ok = True
    by_template: dict[str, list[Fix]] = {}
    for fix in fixes():
        by_template.setdefault(fix.template, []).append(fix)
    for rel, items in by_template.items():
        path = templates / rel
        src = path.read_bytes()
        before = _snapshot(src)
        pending = [f for f in items if before[f.sheet][0].get(f.coord) != f.new]
        for f in pending:
            if before[f.sheet][0].get(f.coord) != f.old:
                print(f"[FAIL] {rel} {f.sheet}!{f.coord} 现状 {before[f.sheet][0].get(f.coord)!r} ≠ 预期旧值 {f.old!r}")
                ok = False
        if not pending:
            print(f"[done] {rel}：{len(items)} 格均已是期望值")
            continue
        if not ok:
            continue
        out = patch_bytes(src, pending)
        after = _snapshot(out)
        changed = {
            (title, coord)
            for title in before
            for coord in set(before[title][0]) | set(after[title][0])
            if before[title][0].get(coord) != after[title][0].get(coord)
        }
        expected = {(f.sheet, f.coord) for f in pending}
        wrong = [f for f in pending if after[f.sheet][0].get(f.coord) != f.new]
        merged = [t for t in before if before[t][1] != after[t][1]]
        if changed != expected or wrong or merged:
            print(f"[FAIL] {rel}：多改 {sorted(changed - expected)[:5]} 漏改 {sorted(expected - changed)[:5]} "
                  f"值不对 {[(f.coord, after[f.sheet][0].get(f.coord)) for f in wrong][:5]} merged {merged}")
            ok = False
            continue
        print(f"[ok] {rel}：{len(pending)} 格  sha256 {hashlib.sha256(src).hexdigest()[:12]} → "
              f"{hashlib.sha256(out).hexdigest()[:12]}")
        if apply:
            path.write_bytes(out)
    print("全部判据通过" if ok else "存在失败，未写盘的模板见上")
    return 0 if ok else 1


if __name__ == "__main__":
    args = sys.argv[1:]
    if "--apply" not in args and "--check" not in args:
        print(__doc__)
        raise SystemExit(2)
    # --templates <目录>：在隔离副本上预演 / 修复（缺省 = 本仓库 backend/wp_templates）
    root = Path(args[args.index("--templates") + 1]) if "--templates" in args else TEMPLATES
    raise SystemExit(run(apply="--apply" in args, templates=root))

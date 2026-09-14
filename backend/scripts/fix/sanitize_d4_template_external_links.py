# -*- coding: utf-8 -*-
"""净化 D4 收入底稿权威模板的外链（过 OOXML external_relationships 门）。

G5-1 Phase 5 · D4 canary（spec: d4-revenue-matrix-bidirectional · Task 4）

与 D3/D5/D6/D7 的关键差别：**公式分类**。D4 有 270 个含 `[n]` 的外链公式格，
同时受管 sheet `主营业务收入明细表D4-2` 的 N 列是 shared `=SUM(B:M)` **合法内部公式**，
不得误删。

净化（zip 级）：
  ① 删 17 个 xl/externalLinks/externalLink*.xml + 对应 _rels
  ② [Content_Types].xml 删 Override
  ③ workbook.xml.rels 删 externalLink Relationship
  ④ workbook.xml 删 <externalReferences> + 含 [n] 的 definedName
  ⑤ 全部 worksheet：只去 ``<f>[n]...</f>``（270 处全是无属性 plain f），
     保留 <v>；内部公式（含 shared SUM）逐字节不变

判据（Property 3/4）：
  - 受管 sheet 内**不含 [n]** 的格子逐格 0 diff（含 N 列公式文本）
  - merge 范围不变
  - `validate_ooxml_artifact(..., document_type='xlsx')` PASS
  - `.preclean.bak` 仍 REJECT（门负例）

用法（仓库根，PowerShell）：
  & .venv/Scripts/python.exe backend/scripts/fix/sanitize_d4_template_external_links.py --check
  & .venv/Scripts/python.exe backend/scripts/fix/sanitize_d4_template_external_links.py --apply
"""
from __future__ import annotations

import hashlib
import io
import re
import shutil
import sys
import tempfile
import zipfile
from pathlib import Path

_REPO = Path(__file__).resolve().parents[3]
TEMPLATE = _REPO / "backend" / "wp_templates" / "D" / "D4 收入底稿.xlsx"
_MANAGED_SHEET = "主营业务收入明细表D4-2"

_EXT_LINK_PART_RE = re.compile(r"^xl/externalLinks/externalLink\d+\.xml$")
_EXT_LINK_RELS_RE = re.compile(r"^xl/externalLinks/_rels/externalLink\d+\.xml\.rels$")
_WORKSHEET_RE = re.compile(r"^xl/worksheets/sheet\d+\.xml$")
_BRACKET_RE = re.compile(r"\[\d+\]")


def _strip_external_from_content_types(xml: str) -> str:
    return re.sub(
        r'<Override PartName="/xl/externalLinks/externalLink\d+\.xml"[^>]*/>',
        "",
        xml,
    )


def _strip_external_from_workbook_rels(xml: str) -> str:
    return re.sub(
        r'<Relationship [^>]*Target="externalLinks/externalLink\d+\.xml"[^>]*/>',
        "",
        xml,
    )


def _strip_external_from_workbook(xml: str) -> str:
    xml = re.sub(r"<externalReferences>.*?</externalReferences>", "", xml, flags=re.S)

    def _drop_bracket_dn(m: re.Match) -> str:
        return "" if re.search(r"\[\d+\]", m.group(0)) else m.group(0)

    xml = re.sub(r"<definedName [^>]*>.*?</definedName>", _drop_bracket_dn, xml, flags=re.S)
    return xml


def _neutralize_external_formulas(xml: str) -> str:
    """去掉含 [n] 的 plain <f>...</f>，保留同 <c> 内 <v>。

    实测 270 处全是无属性 plain f；不得匹配 shared 内部公式
    ``<f t="shared" ...>SUM(B12:M12)</f>``。
    """
    return re.sub(r"<f>[^<]*\[\d+\][^<]*</f>", "", xml)


def sanitize_bytes(src: bytes) -> bytes:
    zin = zipfile.ZipFile(io.BytesIO(src))
    out = io.BytesIO()
    zout = zipfile.ZipFile(out, "w", zipfile.ZIP_DEFLATED)
    for info in zin.infolist():
        name = info.filename
        if _EXT_LINK_PART_RE.match(name) or _EXT_LINK_RELS_RE.match(name):
            continue
        data = zin.read(name)
        if name == "[Content_Types].xml":
            data = _strip_external_from_content_types(data.decode("utf-8")).encode("utf-8")
        elif name == "xl/_rels/workbook.xml.rels":
            data = _strip_external_from_workbook_rels(data.decode("utf-8")).encode("utf-8")
        elif name == "xl/workbook.xml":
            data = _strip_external_from_workbook(data.decode("utf-8")).encode("utf-8")
        elif _WORKSHEET_RE.match(name):
            data = _neutralize_external_formulas(data.decode("utf-8")).encode("utf-8")
        # 保留源 ZipInfo 元数据，避免 date_time=now 导致 sha256 每次漂移（Task 5 哨兵）
        new_info = zipfile.ZipInfo(filename=name, date_time=info.date_time)
        new_info.compress_type = zipfile.ZIP_DEFLATED
        new_info.external_attr = info.external_attr
        zout.writestr(new_info, data)
    zout.close()
    zin.close()
    return out.getvalue()


def _snapshot_managed(path_or_bytes) -> tuple[dict[str, str], list[str], dict[str, str]]:
    """返回 (全非空格快照, merge, N 列公式快照)。"""
    import openpyxl
    from openpyxl.utils import get_column_letter

    src = (
        path_or_bytes
        if isinstance(path_or_bytes, (str, Path))
        else io.BytesIO(path_or_bytes)
    )
    wb = openpyxl.load_workbook(src, data_only=False)
    ws = wb[_MANAGED_SHEET]
    snap: dict[str, str] = {}
    n_formulas: dict[str, str] = {}
    for r in range(1, 33):
        for c in range(1, 23):
            cell = ws.cell(r, c)
            v = cell.value
            if v is None:
                continue
            coord = f"{get_column_letter(c)}{r}"
            snap[coord] = str(v)
            if c == 14 and isinstance(v, str) and v.startswith("="):
                n_formulas[coord] = v
    merged = sorted(str(m) for m in ws.merged_cells.ranges)
    wb.close()
    return snap, merged, n_formulas


def _internal_cell_diffs(
    before: dict[str, str], after: dict[str, str]
) -> list[tuple[str, str | None, str | None]]:
    """只比较净化前不含 [n] 的格子（Property 3）。含 [n] 的外链公式允许被中和。"""
    diffs: list[tuple[str, str | None, str | None]] = []
    for key in sorted(set(before) | set(after)):
        b = before.get(key)
        a = after.get(key)
        if b is not None and _BRACKET_RE.search(b):
            continue  # 外链公式格：允许变
        if b != a:
            diffs.append((key, b, a))
    return diffs


def _gate_pass(path: Path) -> tuple[bool, str]:
    sys.path.insert(0, str(_REPO / "backend"))
    from app.services.workpaper_sync.ooxml_security import validate_ooxml_artifact

    try:
        validate_ooxml_artifact(path, document_type="xlsx")
        return True, "PASS"
    except Exception as exc:  # noqa: BLE001 — 门负例要吃掉并回报
        return False, f"REJECT:{type(exc).__name__}:{exc}"


def main() -> int:
    apply = "--apply" in sys.argv[1:]
    src = TEMPLATE.read_bytes()
    before_snap, before_merged, before_n = _snapshot_managed(TEMPLATE)
    before_sha = hashlib.sha256(src).hexdigest()

    out = sanitize_bytes(src)
    after_sha = hashlib.sha256(out).hexdigest()
    after_snap, after_merged, after_n = _snapshot_managed(out)
    diffs = _internal_cell_diffs(before_snap, after_snap)

    with tempfile.TemporaryDirectory() as tmp:
        clean_path = Path(tmp) / "d4-clean.xlsx"
        clean_path.write_bytes(out)
        gate_ok, gate_msg = _gate_pass(clean_path)

    n_same = before_n == after_n
    n_data = {
        k: v
        for k, v in after_n.items()
        if k.startswith("N") and k[1:].isdigit() and 12 <= int(k[1:]) <= 23
    }
    n_sum_ok = all(
        v.startswith("=SUM(B") and ":M" in v for v in n_data.values()
    ) and len(n_data) == 12

    print(f"before sha256={before_sha}")
    print(f"after  sha256={after_sha}")
    print(f"ooxml gate (clean): {gate_msg}")
    print(
        f"managed internal-cell diffs={len(diffs)} "
        f"(bracket cells excluded from compare)"
    )
    for k, b, a in diffs[:30]:
        print(f"  {k}: {b!r} -> {a!r}")
    print(f"managed merged same={before_merged == after_merged}")
    print(f"N-column formulas unchanged={n_same} count={len(after_n)}")
    print(f"N12:N23 still SUM(B:M)={n_sum_ok} sample={list(n_data.items())[:2]}")

    ok = (
        gate_ok
        and len(diffs) == 0
        and before_merged == after_merged
        and n_same
        and n_sum_ok
    )
    if not ok:
        print("[FAIL] 净化不满足判据 —— 不写盘")
        return 1

    if apply:
        bak = TEMPLATE.with_suffix(TEMPLATE.suffix + ".preclean.bak")
        if not bak.exists():
            shutil.copy2(TEMPLATE, bak)
            print(f"[apply] 备份门负例 → {bak.name}")
        TEMPLATE.write_bytes(out)
        bak_ok, bak_msg = _gate_pass(bak)
        print(f"[apply] 已净化写盘 {TEMPLATE.name}  new_sha256={after_sha}")
        print(f"[apply] .preclean.bak gate: {bak_msg} (期望 REJECT, ok={not bak_ok})")
        if bak_ok:
            print("[FAIL] 门负例不应 PASS")
            return 1
    else:
        print(
            "[check] 判据全过（门 PASS + 受管内非[n]格 0 diff + merge 不变 + N 列公式不变）；"
            "--apply 才写盘"
        )
    return 0


if __name__ == "__main__":
    raise SystemExit(main())

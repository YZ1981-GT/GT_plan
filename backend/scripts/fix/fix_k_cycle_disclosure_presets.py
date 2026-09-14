#!/usr/bin/env python
"""补齐 K 循环 26 张披露 sheet 的公式预设（幂等）。

**为什么 K 不能照抄 G 的脚本**：`fix_g_cycle_disclosure_presets.py` 把披露 sheet 名
硬编码成两个常量（`附注披露信息（上市公司）` / `附注披露信息（国企）`）—— 那在 G 循环
成立是因为 G 的括号写法统一。**K 循环有 6 种写法**（openpyxl 实测）::

    K1   '附注披露信息(上市公司）'   ← 前半角后全角
    K3   '附注披露信息(上市公司)'    ← 全半角
    K5   '附注披露信息（上市公司)'   ← 前全角后半角
    K10  '附注披露信息（上市公司）'  ← 全全角
    K6   '附注披露信息(国企）'
    K7   '附注披露信息（国有企业）'  ← 是「国有企业」不是「国企」

硬编码任何一种都会让另外 5 种全部失配、块永不命中。故本脚本**运行时 openpyxl 直读**
`backend/wp_templates/K/*.xlsx` 取真名（Property 28 / Requirement 6.7）。

**块设计**（对齐 G/D 循环既有范式，见 `prefill_formula_mapping.json` 里 57 个既有披露块）：

- 余额类（K1~K7）：``期末账面余额`` = ``TB(原值,'期末余额')``；有备抵的另加
  ``期末减值准备`` 与 ``期末账面价值``（= 余额 − 备抵）；``期初账面余额``；
  ``上年账面价值`` = ``PREV(本披露页,'期末账面价值')``；``审定表核对`` = ``WP(审定表,'审定数')``
- 损益类（K8~K13）：``本期发生额`` = ``TB(码,'本期发生额')``；``上期发生额`` =
  ``PREV(审定表,'未审数')``；``审定表核对`` = ``WP(审定表,'审定数')``
- K4（宁缺勿造，`has_account=False`）：**不建任何 TB cell**，只建审定表勾稽 ——
  它的科目在 `account_chart` 两侧都零命中，造 `TB()` 就是造假数据
- K0：源 xlsx 无披露 sheet，天然跳过

**环检测**（Property 27）：披露块只**单向**引用审定表；本脚本不给明细表块加任何
指向审定表的 `WP()`，也不给披露块加指向明细表的引用（避免
审定表→明细表→披露→审定表 的三角环）。

Usage::

    python backend/scripts/fix/fix_k_cycle_disclosure_presets.py            # dry-run
    python backend/scripts/fix/fix_k_cycle_disclosure_presets.py --check    # exit 1 = 有欠账
    python backend/scripts/fix/fix_k_cycle_disclosure_presets.py --apply

spec: .kiro/specs/k-cycle-extraction-formula-and-disclosure-closure/
      Task 12 / Requirements 7.1~7.6 / Property 26, 27, 28
"""

from __future__ import annotations

import argparse
import json
import re
import sys
from pathlib import Path

_BACKEND = Path(__file__).resolve().parents[2]
sys.path.insert(0, str(_BACKEND))

MAPPING_PATH = _BACKEND / "data" / "prefill_formula_mapping.json"
TPL_DIR = _BACKEND / "wp_templates" / "K"

#: 披露 sheet 的识别式（源 xlsx 真名一律以此开头，括号写法各异）
DISCLOSURE_RE = re.compile(r"^附注披露信息")

#: 变体判定：sheet 名里含「上市」= listed，含「国企」或「国有企业」= soe
LISTED_RE = re.compile(r"上市")
SOE_RE = re.compile(r"国企|国有企业")

#: 审定表 sheet 的识别式（K4/K5/K6 的真名带空格，故用正则而非字面量）
ADJ_SHEET_RE = re.compile(r"^审定表\s*K\d+-1$")

#: 🔴 无预设登记表：sheet → 理由（≥15 字）。
#:
#: Requirement 7.2 要求「无预设的 sheet 必须在登记表内且有理由」——
#: 目的是禁止**沉默的缺口**。K4 是唯一进这张表的循环。
NO_PRESET_REGISTRY: dict[str, str] = {
    "K4": (
        "K4 其他流动负债走宁缺勿造：BS-053 四变体 formula 全 NULL，且 account_chart "
        "按名（其他流动负债）按码（2301）两侧都零命中 ⇒ 该科目在实务中是报表行、由多个"
        "明细按性质归集，四表侧无从取数。造 TB() 等于造假数据，故披露块只留审定表勾稽。"
    ),
}

#: 白名单：本来就不该有取数预设的 sheet（Requirement 7.6 要求配非空自检）
WHITELIST_RE = re.compile(r"^(底稿目录|GT_Custom|会计提示)$|程序表\s*K\d+A?$|K\d+A$")


# ─────────────────────────────────────────────────────────────────────────────
# 源 xlsx 事实（openpyxl 直读，禁硬编码）
# ─────────────────────────────────────────────────────────────────────────────


def read_source_sheets() -> dict[str, list[str]]:
    """``{wp_code: [sheet 名, ...]}``（运行时权威 = `backend/wp_templates/K/`）。

    🔴 取不到时**抛异常**而非返回空 dict —— 静默返空会让本脚本报「无需变更」，
    把「模板读不到」伪装成「已经收敛」（fail-open 掩盖接线错误的经典形态）。
    """
    from openpyxl import load_workbook

    if not TPL_DIR.is_dir():
        raise RuntimeError(f"源模板目录不存在：{TPL_DIR}")
    out: dict[str, list[str]] = {}
    for f in sorted(TPL_DIR.glob("*.xlsx")):
        if f.name.startswith("~$"):  # WPS/Excel 锁文件
            continue
        wp = f.name.split()[0].upper()
        wb = load_workbook(f, read_only=True, data_only=True)
        try:
            out[wp] = list(wb.sheetnames)
        finally:
            wb.close()
    if not out:
        raise RuntimeError(f"{TPL_DIR} 下没读到任何 xlsx —— 判据源为空，拒绝继续")
    return out


def disclosure_sheets(sheets: dict[str, list[str]]) -> list[tuple[str, str, str]]:
    """``[(wp_code, sheet 真名, 'listed'|'soe'), ...]``。"""
    out: list[tuple[str, str, str]] = []
    for wp, names in sheets.items():
        for n in names:
            if not DISCLOSURE_RE.match(n):
                continue
            if LISTED_RE.search(n):
                out.append((wp, n, "listed"))
            elif SOE_RE.search(n):
                out.append((wp, n, "soe"))
            # 既不含「上市」也不含「国企/国有企业」⇒ 变体判不出，留给守卫报错
    return out


def adjudication_sheet(sheets: dict[str, list[str]], wp: str) -> str | None:
    """该循环审定表 sheet 的**真名**（K4/K5/K6 带空格）。"""
    for n in sheets.get(wp, []):
        if ADJ_SHEET_RE.match(n):
            return n
    return None


# ─────────────────────────────────────────────────────────────────────────────
# 建块
# ─────────────────────────────────────────────────────────────────────────────


def _variant_label(variant: str) -> str:
    return "上市" if variant == "listed" else "国企"


#: 🔴 披露块的 cell_ref **必须带变体后缀**。
#:
#: `formula_management.preset_library.convert_prefill_presets()` 按
#: ``(page_key, target_cell)`` **二元组**去重，而 ``page_key = f"workpaper:{wp_code}"``
#: **不含 sheet** ⇒ 同一循环的 listed / soe 两个披露块若用相同 cell_ref，
#: 第二个会被**静默丢弃**（不报错、不警告）。
#:
#: 这在平台上是既有系统性缺口：实测 D1 / D2 / G14 的两变体披露块 cell_ref
#: **完全撞键**，只是它们没有对应守卫所以没人发现；N2 因两变体披露内容本就不同
#: 而恰好躲过。K1 有 `test_k1_no_dedup_collision_within_page_key` 守着，
#: 建块时立刻打红 —— 这条守卫救了一次。
#:
#: 治法：给 cell_ref 加变体后缀。这些是**新建锚点**、无既有绑定，
#: 故加后缀不会破坏 `prefill_anchor_map`（它本就用三元组）或用户已录入的公式覆盖。
_VARIANT_SUFFIX: dict[str, str] = {"listed": "_上市", "soe": "_国企"}


def build_block(wp: str, sheet: str, variant: str, adj_sheet: str | None) -> dict:
    """构造一个披露块（对齐 G/D 既有范式）。"""
    from app.services.four_table.k_cycle_specs import K_CYCLE_SPECS

    spec = K_CYCLE_SPECS[wp]
    name = spec.account_name
    label = _variant_label(variant)
    sfx = _VARIANT_SUFFIX[variant]
    gross = spec.fallback_standard
    provisions = [c for c in spec.fallback_provision if c]
    if wp == "K6":
        provisions = provisions or ["1482"]

    cells: list[dict] = []

    if gross and spec.is_pl:
        cells.append(
            {
                "cell_ref": f"{name}_本期{sfx}",
                "formula": f"=TB('{gross}','本期发生额')",
                "formula_type": "TB",
                "description": f"{label}披露表{name}本期发生额（损益类无期末余额口径）",
            }
        )
        if adj_sheet:
            cells.append(
                {
                    "cell_ref": f"{name}_上期{sfx}",
                    "formula": f"=PREV('{wp}','{adj_sheet}','未审数')",
                    "formula_type": "PREV",
                    "description": (
                        f"{label}披露表{name}上期发生额（取上年同底稿审定表的「未审数」格）"
                    ),
                }
            )
    elif gross:
        cells.append(
            {
                "cell_ref": f"期末账面余额{sfx}",
                "formula": f"=TB('{gross}','期末余额')",
                "formula_type": "TB",
                "description": f"{label}披露表{name}期末账面余额（科目 {gross}）",
            }
        )
        if provisions:
            prov_expr = "-".join(f"TB('{c}','期末余额')" for c in provisions)
            cells.append(
                {
                    "cell_ref": f"期末减值准备{sfx}",
                    "formula": "=" + "+".join(f"TB('{c}','期末余额')" for c in provisions),
                    "formula_type": "TB",
                    "description": (
                        f"{label}披露表{name}期末减值准备（备抵码 {', '.join(provisions)}）"
                    ),
                }
            )
            cells.append(
                {
                    "cell_ref": f"期末账面价值{sfx}",
                    "formula": f"=TB('{gross}','期末余额')-{prov_expr}",
                    "formula_type": "TB",
                    "description": f"{label}披露表{name}期末账面价值 = 账面余额 − 减值准备",
                }
            )
        cells.append(
            {
                "cell_ref": f"期初账面余额{sfx}",
                "formula": f"=TB('{gross}','期初余额')",
                "formula_type": "TB",
                "description": f"{label}披露表{name}期初账面余额（科目 {gross}）",
            }
        )
        cells.append(
            {
                "cell_ref": f"上年账面价值{sfx}",
                # PREV 第三参指向**本披露页**的账面价值格（含同一变体后缀）
                "formula": f"=PREV('{wp}','{sheet}','期末账面价值{sfx}')",
                "formula_type": "PREV",
                "description": f"上年{name}账面价值（取上年同披露页，{label}变体）",
            }
        )

    # 审定表勾稽（**单向**：披露 → 审定表。不反向，且不引明细表 ——
    # 否则与既有「审定表 → 明细表」合成 审定表→明细表→披露→审定表 三角环）
    if adj_sheet:
        cells.append(
            {
                "cell_ref": f"审定表核对{sfx}",
                "formula": f"=WP('{wp}','{adj_sheet}','审定数')",
                "formula_type": "WP",
                "description": f"与{adj_sheet}的审定数勾稽（不参与披露取数）",
            }
        )

    block: dict = {
        "wp_code": wp,
        "wp_name": f"{name}附注披露（{label}）",
        "sheet": sheet,
        "account_codes": [c for c in ([gross] if gross else []) + provisions],
        "cells": cells,
    }
    if wp in NO_PRESET_REGISTRY:
        block["_no_tb_preset_reason"] = NO_PRESET_REGISTRY[wp]
    return block


# ─────────────────────────────────────────────────────────────────────────────
# 计划 / 应用
# ─────────────────────────────────────────────────────────────────────────────


def plan(data: dict, sheets: dict[str, list[str]]) -> tuple[list[dict], list[str]]:
    """返回 ``(要新增的块, 变更日志)``。

    幂等口径：
    - 已存在且 cell_ref **已带变体后缀**的块 → 不动
    - 已存在但 cell_ref **缺后缀**的块 → 就地重建（迁移；本脚本首版漏了后缀）
    - 不存在 → 新建
    """
    by_key = {
        (str(b.get("wp_code", "")).upper(), str(b.get("sheet", ""))): b
        for b in data["mappings"]
    }
    additions: list[dict] = []
    changes: list[str] = []
    for wp, sheet, variant in sorted(disclosure_sheets(sheets)):
        adj = adjudication_sheet(sheets, wp)
        blk = build_block(wp, sheet, variant, adj)
        old = by_key.get((wp, sheet))
        if old is None:
            additions.append(blk)
            changes.append(
                f"ADD disclosure block: {wp} | {sheet!r} | {len(blk['cells'])} cells "
                f"| accounts={blk['account_codes']}"
                + ("" if adj else "  ⚠ 未找到审定表 sheet，缺勾稽 cell")
            )
            continue
        sfx = _VARIANT_SUFFIX[variant]
        old_refs = [str(c.get("cell_ref")) for c in old.get("cells") or []]
        if old_refs and all(r.endswith(sfx) for r in old_refs):
            continue  # 已是目标形态
        # 缺变体后缀 ⇒ 就地重建（cell_ref 撞键会让 convert_prefill_presets 静默丢条目）
        old["wp_name"] = blk["wp_name"]
        old["account_codes"] = blk["account_codes"]
        old["cells"] = blk["cells"]
        if "_no_tb_preset_reason" in blk:
            old["_no_tb_preset_reason"] = blk["_no_tb_preset_reason"]
        changes.append(
            f"MIGRATE cell_ref 加变体后缀 {sfx}: {wp} | {sheet!r} | "
            f"{len(blk['cells'])} cells（原 {old_refs[:3]}…）"
        )
    return additions, changes


def _round_trip_guard(raw: str) -> None:
    """`json.dumps` 不能逐字复现原文即 exit 2（拒绝写盘）。

    该文件 429KB / 281 块、被 5 个 spec 共享。序列化形态不一致会让一次写盘
    重排全文、把并发会话的成果全部卷进 diff。
    """
    if json.dumps(json.loads(raw), ensure_ascii=False, indent=2) + "\n" == raw:
        return
    print(
        "round-trip 自检失败：json.dumps(indent=2)+换行 无法逐字复现原文，已拒绝写盘。",
        file=sys.stderr,
    )
    sys.exit(2)


def _utf8_stdout() -> None:
    for s in (sys.stdout, sys.stderr):
        rc = getattr(s, "reconfigure", None)
        if rc is not None:
            try:
                rc(encoding="utf-8", errors="replace")
            except Exception:  # noqa: BLE001
                pass


def main() -> int:
    _utf8_stdout()
    ap = argparse.ArgumentParser(description="K 循环披露块预设补齐（幂等）")
    ap.add_argument("--apply", action="store_true", help="执行并写盘")
    ap.add_argument("--check", action="store_true", help="有欠账则 exit 1")
    args = ap.parse_args()

    raw = MAPPING_PATH.read_text(encoding="utf-8")
    _round_trip_guard(raw)
    data = json.loads(raw)

    sheets = read_source_sheets()
    discs = disclosure_sheets(sheets)
    print(f"源 xlsx 实测披露 sheet {len(discs)} 张（{len(sheets)} 个工作簿）")

    additions, changes = plan(data, sheets)
    if not changes:
        print("无需变更（26 张披露 sheet 已全部有预设块）")
        return 0

    print(f"{'[DRY-RUN] ' if not args.apply else ''}变更清单（{len(changes)} 项）:")
    for c in changes:
        print(f"  - {c}")

    if args.apply:
        data["mappings"].extend(additions)
        MAPPING_PATH.write_text(
            json.dumps(data, ensure_ascii=False, indent=2) + "\n", encoding="utf-8"
        )
        print(f"\n已写盘：{MAPPING_PATH}")
        return 0
    if args.check:
        print(f"\n有 {len(changes)} 项欠账")
        return 1
    print("\nDRY-RUN 模式，使用 --apply 执行")
    return 0


if __name__ == "__main__":
    raise SystemExit(main())

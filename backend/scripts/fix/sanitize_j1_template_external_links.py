# -*- coding: utf-8 -*-
"""净化 J1 应付职工薪酬权威模板的孤儿外链（过 OOXML external_relationships 门）。

spec: `j-cycle-sync-foundation-and-first-canary`（22e 首版发布被 `ooxml_security_rejected` 挡下后）
范式：照 `sanitize_d7_template_external_links.py`（D3~D7 同款，对齐 DEC-11 / B60 净化范式）——
`backend/wp_templates/` 运行时只读，模板真要升级必须显式净化 + 留 `.preclean.bak` 作门负例 +
重算 sentinel 重发布。

净化对象（2026-10-01 逐格现查，全部是**断链孤儿**，不影响受管 sheet `计提情况检查表J1-6`）：
  ① 删 2 个 `xl/externalLinks/externalLink*.xml` + 2 个 `_rels/*.rels`
     （目标是旧作者本机路径 `file:///D:\\Documents and Settings\\…` 与 `D:\\5.底稿模板\\…\\L1 应付职工薪酬.xlsx`）
  ② `[Content_Types].xml` 删对应 Override；③ `xl/_rels/workbook.xml.rels` 删对应 Relationship
  ④ `xl/workbook.xml` 删 `<externalReferences>` 块 + 删含 `[n]` 的 defined name（现算 1 个）
  ⑤ 含 `[n]` 的公式格（现算 6 格，**全在 hidden 串册 sheet `应付职工薪酬实质性程序表 L1A-原`** 的
     R3/R4 表头 `=[2]底稿目录!A2…A7`）去 `<f>` 保 `<v>` 缓存值（可见内容不变）
     —— 与 D7 不同处：sheet 文件名**现算**定位（按 `<f>` 含 `[n]`），不写死 sheetN.xml

用法（仓库根）：
  & .venv/Scripts/python.exe backend/scripts/fix/sanitize_j1_template_external_links.py          # 预演
  & .venv/Scripts/python.exe backend/scripts/fix/sanitize_j1_template_external_links.py --apply  # 写盘 + .bak

🔴 判成败用 openpyxl 打开 + 受管 sheet 逐格 diff + 门判据，不看退出码。
"""
from __future__ import annotations

import hashlib
import io
import re
import shutil
import sys
import zipfile
from pathlib import Path

_REPO = Path(__file__).resolve().parents[3]
TEMPLATE = _REPO / "backend" / "wp_templates" / "J" / "J1 应付职工薪酬.xlsx"
_MANAGED_SHEET = "计提情况检查表J1-6"

_EXT_LINK_PART_RE = re.compile(r"^xl/externalLinks/externalLink\d+\.xml$")
_EXT_LINK_RELS_RE = re.compile(r"^xl/externalLinks/_rels/externalLink\d+\.xml\.rels$")
_EXT_FORMULA_RE = re.compile(r"<f>[^<]*\[\d+\][^<]*</f>")


def _strip_content_types(xml: str) -> str:
    return re.sub(r'<Override PartName="/xl/externalLinks/externalLink\d+\.xml"[^>]*/>', "", xml)


def _strip_workbook_rels(xml: str) -> str:
    return re.sub(r'<Relationship [^>]*Target="externalLinks/externalLink\d+\.xml"[^>]*/>', "", xml)


def _strip_workbook(xml: str) -> str:
    xml = re.sub(r"<externalReferences>.*?</externalReferences>", "", xml, flags=re.S)

    def _drop(m: re.Match) -> str:
        return "" if re.search(r"\[\d+\]", m.group(0)) else m.group(0)

    return re.sub(r"<definedName [^>]*>.*?</definedName>", _drop, xml, flags=re.S)


def sanitize_bytes(src: bytes) -> tuple[bytes, dict[str, int]]:
    zin = zipfile.ZipFile(io.BytesIO(src))
    out = io.BytesIO()
    zout = zipfile.ZipFile(out, "w", zipfile.ZIP_DEFLATED)
    stats = {"dropped_parts": 0, "neutralized_formulas": 0}
    for info in zin.infolist():
        name = info.filename
        if _EXT_LINK_PART_RE.match(name) or _EXT_LINK_RELS_RE.match(name):
            stats["dropped_parts"] += 1
            continue
        data = zin.read(name)
        if name == "[Content_Types].xml":
            data = _strip_content_types(data.decode("utf-8")).encode("utf-8")
        elif name == "xl/_rels/workbook.xml.rels":
            data = _strip_workbook_rels(data.decode("utf-8")).encode("utf-8")
        elif name == "xl/workbook.xml":
            data = _strip_workbook(data.decode("utf-8")).encode("utf-8")
        elif name.startswith("xl/worksheets/sheet") and name.endswith(".xml"):
            text = data.decode("utf-8")
            hits = len(_EXT_FORMULA_RE.findall(text))
            if hits:
                stats["neutralized_formulas"] += hits
                data = _EXT_FORMULA_RE.sub("", text).encode("utf-8")
        zout.writestr(name, data)
    zout.close()
    zin.close()
    return out.getvalue(), stats


def _snapshot_managed(data: bytes) -> tuple[dict[str, str], list[str]]:
    import openpyxl

    wb = openpyxl.load_workbook(io.BytesIO(data), data_only=False)
    ws = wb[_MANAGED_SHEET]
    snap = {
        c.coordinate: str(c.value)
        for row in ws.iter_rows(min_row=1, max_row=ws.max_row, max_col=ws.max_column)
        for c in row
        if c.value is not None
    }
    merged = sorted(str(m) for m in ws.merged_cells.ranges)
    wb.close()
    return snap, merged


def _external_rels_count(data: bytes) -> int:
    z = zipfile.ZipFile(io.BytesIO(data))
    n = sum(
        1
        for name in z.namelist()
        if name.endswith(".rels") and 'TargetMode="External"' in z.read(name).decode("utf-8", "replace")
    )
    z.close()
    return n


def main() -> int:
    apply = "--apply" in sys.argv[1:]
    src = TEMPLATE.read_bytes()
    out, stats = sanitize_bytes(src)
    before_snap, before_merged = _snapshot_managed(src)
    after_snap, after_merged = _snapshot_managed(out)
    diffs = [k for k in sorted(set(before_snap) | set(after_snap)) if before_snap.get(k) != after_snap.get(k)]
    ext_before, ext_after = _external_rels_count(src), _external_rels_count(out)

    print(f"before sha256={hashlib.sha256(src).hexdigest()} size={len(src)}")
    print(f"after  sha256={hashlib.sha256(out).hexdigest()} size={len(out)}")
    print(f"stats={stats}")
    print(f"external .rels: before={ext_before} after={ext_after} -> {'PASS' if ext_after == 0 else 'REJECT'}")
    print(f"managed cells before={len(before_snap)} after={len(after_snap)} diffs={len(diffs)} {diffs[:10]}")
    print(f"managed merged same={before_merged == after_merged}")

    ok = (
        ext_before > 0
        and ext_after == 0
        and not diffs
        and before_merged == after_merged
        and stats["dropped_parts"] == 4
        and stats["neutralized_formulas"] == 6
    )
    if not ok:
        print("[FAIL] 判据不满足（门 / 受管 sheet diff / merge / 现算计数 4·6 不符）—— 不写盘")
        return 1
    if apply:
        bak = TEMPLATE.with_suffix(TEMPLATE.suffix + ".preclean.bak")
        if not bak.exists():
            shutil.copy2(TEMPLATE, bak)
            print(f"[apply] 备份门负例 -> {bak.name}")
        TEMPLATE.write_bytes(out)
        print("[apply] 已净化写盘")
    else:
        print("[check] 判据全过；--apply 才写盘")
    return 0


if __name__ == "__main__":
    raise SystemExit(main())

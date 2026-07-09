"""Analyze S35 再融资审核特项底稿 structure.

Reads 5 S35 xlsx files and outputs:
- Sheet names (核查程序表 + 明细核查子表 S35-x-1)
- Formula columns (columns with formulas)
- Judgment/check columns (核查判断列)
- Column headers
- Structure summary for bundle TabDef generation

Read-only analysis; writes nothing to source templates.
"""
from __future__ import annotations

import sys
from pathlib import Path

import openpyxl

ROOT = Path(
    r"d:\GT_plan\基础数据\致同通用审计程序及底稿模板（2025年修订）"
    r"\1.致同审计程序及底稿模板（2025年）\6.特定项目程序（S）"
    r"\S35 再融资审核特项底稿"
)


def analyze_columns(ws, max_header_row: int = 3):
    """Identify column headers, formula columns, and judgment columns."""
    headers = []
    formula_cols = set()
    judgment_cols = set()

    # Read header rows
    for row_idx in range(1, max_header_row + 1):
        row_headers = []
        for col_idx in range(1, (ws.max_column or 0) + 1):
            cell = ws.cell(row=row_idx, column=col_idx)
            v = cell.value
            if v is not None:
                row_headers.append((col_idx, str(v).strip()))
        if row_headers:
            headers.append((row_idx, row_headers))

    # Scan data rows for formulas and judgment keywords
    judgment_keywords = ['核查', '判断', '结论', '是否', '审核', '意见', '说明', '备注']
    for row in ws.iter_rows(min_row=1, max_row=min(ws.max_row or 0, 50)):
        for cell in row:
            v = cell.value
            if isinstance(v, str):
                if v.startswith('='):
                    formula_cols.add(cell.column)
                # Check header cells for judgment keywords
                if cell.row <= max_header_row:
                    for kw in judgment_keywords:
                        if kw in v:
                            judgment_cols.add((cell.column, v))
                            break

    return headers, formula_cols, judgment_cols


def main() -> None:
    files = sorted(ROOT.glob("S35-*.xlsx"), key=lambda p: p.name)
    print(f"# S35 再融资审核特项底稿结构分析\n")
    print(f"发现 {len(files)} 个工作簿文件\n")
    print("=" * 80)

    all_tabs = []

    for p in files:
        print(f"\n## {p.name}")
        print("-" * 60)
        try:
            wb = openpyxl.load_workbook(p, read_only=False, data_only=False)
        except Exception as e:
            print(f"  <ERROR: {e}>")
            continue

        wp_code = p.stem.split(' ')[0]  # e.g. 'S35-1'
        label = p.stem.split(' ', 1)[1] if ' ' in p.stem else ''
        sub_sheets = []

        print(f"  wp_code: {wp_code}")
        print(f"  label: {label}")
        print(f"  sheets: {wb.sheetnames}")
        print()

        for idx, name in enumerate(wb.sheetnames):
            ws = wb[name]
            rows = ws.max_row or 0
            cols = ws.max_column or 0
            merged = len(ws.merged_cells.ranges)

            print(f"  ### Sheet [{idx}]: \"{name}\"")
            print(f"      size: {rows}行 x {cols}列, merged={merged}")

            headers, formula_cols, judgment_cols = analyze_columns(ws)

            # Print header rows
            for row_idx, row_headers in headers:
                print(f"      header_row_{row_idx}: {row_headers}")

            if formula_cols:
                print(f"      formula_cols: {sorted(formula_cols)}")
            if judgment_cols:
                print(f"      judgment_cols: {sorted(judgment_cols, key=lambda x: x[0])}")

            # Determine if this is a program sheet or a sub-check table
            is_sub = '-1' in name or '明细' in name or '核查表' in name
            if idx > 0:
                sub_sheets.append(name)
                print(f"      type: 明细核查子表")
            else:
                print(f"      type: 核查程序表")
            print()

        # Build tab def
        tab_info = {
            'id': wp_code,
            'label': label,
            'wpCode': wp_code,
            'subSheets': sub_sheets if sub_sheets else None,
            'sheetNames': wb.sheetnames,
        }
        all_tabs.append(tab_info)
        wb.close()

    # Summary
    print("\n" + "=" * 80)
    print("\n# 汇总：5 核查底稿 Tab 标签 + 子表映射表\n")
    print("```typescript")
    print("// S35 TabDef 配置")
    print("export const S35_TAB_DEFS: TabDef[] = [")
    for t in all_tabs:
        sub = f", subSheets: {t['subSheets']}" if t['subSheets'] else ""
        print(f"  {{ id: '{t['id']}', label: '{t['label']}', wpCode: '{t['wpCode']}'{sub} }},")
    print("]")
    print("```")
    print()
    print("# 子表映射表")
    print()
    for t in all_tabs:
        if t['subSheets']:
            for s in t['subSheets']:
                print(f"  {t['wpCode']} -> {s} (明细核查子表)")
        else:
            print(f"  {t['wpCode']} -> (无子表)")


if __name__ == "__main__":
    sys.stdout.reconfigure(encoding="utf-8")
    main()

# -*- coding: utf-8 -*-
"""D4-3 Wave 1 census: geometry + column↔field mapping digest."""
from __future__ import annotations

import hashlib
import json
import re
from pathlib import Path

from openpyxl import load_workbook
from openpyxl.utils import get_column_letter

REPO = Path(__file__).resolve().parents[3]
TEMPLATE = REPO / "backend" / "wp_templates" / "D" / "D4 收入底稿.xlsx"
OUT = REPO / ".kiro" / "specs" / "d-cycle-sheet-bidirectional-expansion" / "evidence"
OUT.mkdir(parents=True, exist_ok=True)


def main() -> None:
    sha = hashlib.sha256(TEMPLATE.read_bytes()).hexdigest()
    wb = load_workbook(TEMPLATE, data_only=False)
    name = next(n for n in wb.sheetnames if "D4-3" in n and "其他业务" in n)
    ws = wb[name]

    # dump header band
    band: dict[str, list[str | None]] = {}
    for r in range(10, 22):
        row_vals: list[str | None] = []
        for c in range(1, 15):
            v = ws.cell(r, c).value
            row_vals.append(None if v is None else str(v))
        if any(x is not None for x in row_vals):
            band[f"R{r}"] = row_vals

    # detect footer
    footer_row = None
    footer_marker = None
    for r in range(12, 40):
        a = ws.cell(r, 1).value
        if isinstance(a, str) and a.strip() == "合计":
            footer_row = r
            footer_marker = a
            break

    # formulas
    formulas: list[dict[str, str]] = []
    bracket = 0
    for row in ws.iter_rows(min_row=1, max_row=40, max_col=14):
        for cell in row:
            v = cell.value
            if isinstance(v, str) and v.startswith("="):
                formulas.append({"cell": cell.coordinate, "f": v})
                if re.search(r"\[\d+\]", v):
                    bracket += 1

    # Proposed mapping (HTML StoredOtherRow vs Excel)
    # Decision: D/I 重分类 → formula_mask / not in store (HTML has no field); write 0 if create
    field_map = [
        {"col": "A", "r12": "项目", "mode": "editable", "store_key": "item", "value_type": "text", "contract": True},
        {"col": "B", "r12": "本期未审数", "mode": "editable", "store_key": "currentUnadjusted", "value_type": "amount", "contract": True},
        {"col": "C", "r12": "账项调整", "mode": "editable", "store_key": "currentAdjustment", "value_type": "amount", "contract": True},
        {"col": "D", "r12": "重分类调整", "mode": "mask_or_zero", "store_key": None, "value_type": "amount", "contract": False, "note": "HTML store 无此字段；不进契约，materialize 写 0 / extract 丢弃"},
        {"col": "E", "r12": "本期审定数", "mode": "formula", "store_key": None, "value_type": "amount", "contract": False},
        {"col": "F", "r12": "结构比", "mode": "formula", "store_key": None, "value_type": "ratio", "contract": False},
        {"col": "G", "r12": "上期未审数", "mode": "editable", "store_key": "priorUnadjusted", "value_type": "amount", "contract": True},
        {"col": "H", "r12": "账项调整", "mode": "editable", "store_key": "priorAdjustment", "value_type": "amount", "contract": True},
        {"col": "I", "r12": "重分类调整", "mode": "mask_or_zero", "store_key": None, "value_type": "amount", "contract": False, "note": "同 D"},
        {"col": "J", "r12": "上期审定额", "mode": "formula", "store_key": None, "value_type": "amount", "contract": False},
        {"col": "K", "r12": "结构比", "mode": "formula", "store_key": None, "value_type": "ratio", "contract": False},
        {"col": "L", "r12": "变动额", "mode": "formula", "store_key": None, "value_type": "amount", "contract": False},
        {"col": "M", "r12": "变动率", "mode": "formula", "store_key": None, "value_type": "ratio", "contract": False},
        {"col": "N", "r11": "备注", "mode": "editable", "store_key": "remark", "value_type": "text", "contract": True},
    ]
    contract_fields = [f for f in field_map if f["contract"]]
    digest_payload = {
        "managed_sheet": name,
        "header_rows": [11, 12],
        "first_data_row": 13,
        "last_data_row": (footer_row - 1) if footer_row else None,
        "footer_row": footer_row,
        "footer_marker_exact": footer_marker,
        "contract_fields": [
            {"col": f["col"], "store_key": f["store_key"], "value_type": f["value_type"]}
            for f in contract_fields
        ],
        "mask_cols": [f["col"] for f in field_map if f["mode"] in ("formula", "mask_or_zero")],
        "template_sha256": sha,
    }
    mapping_digest = hashlib.sha256(
        json.dumps(digest_payload, ensure_ascii=False, sort_keys=True, separators=(",", ":")).encode("utf-8")
    ).hexdigest()

    out = {
        "measured_at": "2026-09-12",
        "tool": "openpyxl data_only=False",
        "template_path": "backend/wp_templates/D/D4 收入底稿.xlsx",
        "template_sha256": sha,
        "managed_sheet": name,
        "dimensions": ws.dimensions,
        "header_rows": [11, 12],
        "first_data_row": 13,
        "last_data_row": (footer_row - 1) if footer_row else None,
        "footer_row": footer_row,
        "footer_marker_exact": footer_marker,
        "footer_marker_codepoints": [hex(ord(ch)) for ch in (footer_marker or "")],
        "band_R10_R21": band,
        "formula_count_sample": len(formulas),
        "external_bracket_formulas": bracket,
        "sample_formulas": formulas[:20],
        "merged_ranges": [str(m) for m in ws.merged_cells.ranges],
        "field_map": field_map,
        "contract_field_count": len(contract_fields),
        "store_item_id": "D4-3-rows",
        "frontend_composable": "useD4OtherRevenue.ts / StoredOtherRow",
        "reclass_decision": "D/I 不进契约、不进 HTML store；materialize 填 0；extract 忽略",
        "digest_payload": digest_payload,
        "mapping_digest": mapping_digest,
        "architecture_note": (
            "同一 entry 只能挂一个 adapter（registry）；D4-3 必须扩进 d4.revenue_detail "
            "第二 sheets[] / sheet_key=d43-managed，宿主按 currentSheet 切换 sheetKey"
        ),
    }
    path = OUT / "T01-d43-column-field-mapping.json"
    path.write_text(json.dumps(out, ensure_ascii=False, indent=2) + "\n", encoding="utf-8")
    print("wrote", path)
    print("mapping_digest", mapping_digest)
    print("geometry", digest_payload["first_data_row"], digest_payload["last_data_row"], digest_payload["footer_row"])
    print("contract_fields", len(contract_fields))


if __name__ == "__main__":
    main()

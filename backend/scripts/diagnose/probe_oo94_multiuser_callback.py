"""真实 OnlyOffice 9.4 两用户同 room callback 黑盒探针（spec: workpaper-html-onlyoffice-bidirectional-writeback-closure Task 4）

目的（design.md §多用户 callback 黑盒真值门 / Requirement 4.3、4.10、5.1、5.4、10.2、10.4、10.9）：
在实现 shared room 之前，用**真实 OO 9.4 容器 + 两个独立用户 + 同一 doc_key**取证：

  1. 用户 A 发起 forcesave（Command Service `c=forcesave` 带 `userdata`）时 callback 的
     `status/userdata/users/actions/lastsave/...` 实际内容；
  2. 用户 B 的并发修改是否被聚合进 A 那次 forcesave 产出的 artifact；
  3. status 6（forcesave）与 status 2（最终保存）payload 差异，以及 status 2 是否带 userdata；
  4. `c=drop` 撤销用户 B 后，B 的贡献是否仍留在后续 artifact 中（= 能否证明"已安全 drop"）；
  5. 无 userdata 的 close callback 行为。

为什么**不复用平台生产 callback 端点**：
  - Task 4 硬约束「不接通生产 callback」；
  - 抓原始 payload 需要一个可控收集端点（tasks.md 明确允许「在测试环境起一个独立 callback 收集端点」）；
  - 生产 callback 会把回传文件**覆盖写入真实底稿存储**（`wp_onlyoffice_router.post_sheet_onlyoffice_callback`
    → `storage/projects/{pid}/workpapers/onlyoffice/{wp_code}.xlsx`），探针不应改业务数据。
  故本探针自带 document host + callback collector，文档字节取自 `backend/wp_templates/` 的**真实平台模板**。

用法（Windows，仓库根）：
    python backend/scripts/diagnose/probe_oo94_multiuser_callback.py serve
    python backend/scripts/diagnose/probe_oo94_multiuser_callback.py command --c forcesave --userdata req-001
    python backend/scripts/diagnose/probe_oo94_multiuser_callback.py command --c drop --users bob
    python backend/scripts/diagnose/probe_oo94_multiuser_callback.py state
    python backend/scripts/diagnose/probe_oo94_multiuser_callback.py finalize

约束：
  - 只读容器配置；不改 OO 容器、不改生产代码、不写业务库、不动 `backend/wp_templates/`。
  - evidence 落盘到 `--evidence-dir`（默认 spec evidence 目录），token 类字段一律脱敏成 sha256 前缀。
"""

from __future__ import annotations

import argparse
import hashlib
import io
import json
import os
import shutil
import sys
import threading
import time
import urllib.error
import urllib.request
from datetime import datetime, timezone
from http.server import BaseHTTPRequestHandler, ThreadingHTTPServer
from pathlib import Path
from typing import Any
from urllib.parse import parse_qs, urlsplit, urlunsplit

from jose import jwt

REPO_ROOT = Path(__file__).resolve().parents[3]
DEFAULT_EVIDENCE_DIR = (
    REPO_ROOT
    / ".kiro"
    / "specs"
    / "workpaper-html-onlyoffice-bidirectional-writeback-closure"
    / "evidence"
    / "task4-oo94-multiuser-callback"
)
DEFAULT_SEED = REPO_ROOT / "backend" / "wp_templates" / "A" / "A31 审计标识一览表.xlsx"

# 容器 local.json 实测：token.enable.request.inbox/outbox = true，header = Authorization，inBody = false。
OO_SECRET = os.environ.get("PROBE_OO_SECRET", "onlyoffice-dev-secret-2026")
# 宿主机可达的 OO 地址（.env 的 ONLYOFFICE_URL）。
OO_URL_FROM_HOST = os.environ.get("PROBE_OO_URL", "http://localhost:8080")
# 容器内可达的宿主机地址（.env 的 ONLYOFFICE_CALLBACK_BASE 同款主机名）。
HOST_FROM_CONTAINER = os.environ.get("PROBE_HOST_FOR_OO", "host.docker.internal")

USERS: dict[str, str] = {"alice": "探针用户A", "bob": "探针用户B"}


# ---------------------------------------------------------------------------
# 工具
# ---------------------------------------------------------------------------


def _now() -> str:
    return datetime.now(timezone.utc).isoformat()


def _sha256_bytes(data: bytes) -> str:
    return hashlib.sha256(data).hexdigest()


def _digest_ref(value: str) -> str:
    """把 token/签名类值换成不可逆引用，保证 evidence 可入库。"""
    return f"sha256:{hashlib.sha256(value.encode('utf-8')).hexdigest()[:24]}"


#: callback payload 中属于「URL 且带签名参数」的字段（`url` = 回传文件，`changesurl` = 变更包）。
_URL_FIELDS = ("url", "changesurl", "historyurl")
#: URL query 中必须脱敏的签名参数（`shardkey` 只是 doc key，保留原值便于复核）。
_SIGNED_QUERY_KEYS = {"md5", "token", "signature", "sig"}


def _redact_url(url: str) -> str:
    """OO 回传下载 URL 带 `md5`/`token` 等签名参数 → 值换成 digest 引用后再落盘。"""
    parts = urlsplit(url)
    if not parts.query:
        return url
    redacted_pairs = []
    for raw in parts.query.split("&"):
        if "=" not in raw:
            redacted_pairs.append(raw)
            continue
        key, value = raw.split("=", 1)
        if key.lower() in _SIGNED_QUERY_KEYS:
            redacted_pairs.append(f"{key}={_digest_ref(value)}")
        else:
            redacted_pairs.append(f"{key}={value}")
    return urlunsplit((parts.scheme, parts.netloc, parts.path, "&".join(redacted_pairs), parts.fragment))


def redact_payload(obj: Any) -> Any:
    """递归脱敏 callback payload / JWT claims：URL 签名参数与 `token` 字段。

    evidence 要入库，故**任何签名值都不落原文**；`token`（OO outbox JWT）只留 digest 引用。
    """
    if isinstance(obj, dict):
        out: dict[str, Any] = {}
        for key, value in obj.items():
            if key == "token" and isinstance(value, str):
                out[key] = _digest_ref(value)
            elif key in _URL_FIELDS and isinstance(value, str):
                out[key] = _redact_url(value)
            else:
                out[key] = redact_payload(value)
        return out
    if isinstance(obj, list):
        return [redact_payload(item) for item in obj]
    return obj


def _rewrite_download_host(url: str) -> str:
    """把 OO 自视角 URL（`http://localhost/cache/...` = 容器 80 端口）改写成宿主机可达地址。

    与生产 `wp_onlyoffice_router._rewrite_onlyoffice_download_url` 同口径（scheme+netloc 替换，
    path/query 保留）。此处独立实现是为了让探针不依赖 FastAPI app 启动链；
    生产该函数本身由 `backend/tests/workpaper_bidirectional_writeback/test_task4_callback_download_security_characterization.py`
    直接导入做行为 characterization。
    """
    base = urlsplit(OO_URL_FROM_HOST)
    cur = urlsplit(url)
    if not base.netloc or cur.netloc == base.netloc:
        return url
    return urlunsplit((base.scheme or cur.scheme, base.netloc, cur.path, cur.query, cur.fragment))


def _xlsx_cell_probe(data: bytes, cells: list[str]) -> dict[str, Any]:
    """读回传 xlsx 的指定单元格值（判"B 的贡献是否聚合进 artifact"的唯一判据）。"""
    import openpyxl

    out: dict[str, Any] = {}
    try:
        wb = openpyxl.load_workbook(io.BytesIO(data), data_only=False)
        ws = wb[wb.sheetnames[0]]
        out["sheet_names"] = list(wb.sheetnames)
        out["first_sheet"] = ws.title
        out["cells"] = {ref: ws[ref].value for ref in cells}
        wb.close()
    except Exception as exc:  # noqa: BLE001 — 探针如实记录失败，不吞成"无数据"
        out["error"] = f"{type(exc).__name__}: {exc}"
    return out


def _sign(payload: dict[str, Any]) -> str:
    return jwt.encode(payload, OO_SECRET, algorithm="HS256")


# ---------------------------------------------------------------------------
# 状态（进程内，serve 模式持有）
# ---------------------------------------------------------------------------


class ProbeState:
    def __init__(
        self,
        *,
        evidence_dir: Path,
        seed: Path,
        doc_key: str,
        port: int,
        marker_cells: dict[str, str],
        phase: str,
    ):
        self.evidence_dir = evidence_dir
        self.artifacts_dir = evidence_dir / "artifacts"
        self.seed = seed
        self.doc_key = doc_key
        self.port = port
        self.marker_cells = marker_cells
        self.phase = phase
        #: 剩余多少次 callback 要故意返回 OO 非零 error（实证「返回非零 → OO 是否重发」）。
        self.fail_next = 0
        self.lock = threading.Lock()
        self.seq = 0
        self.callbacks: list[dict[str, Any]] = []
        self.evidence_dir.mkdir(parents=True, exist_ok=True)
        self.artifacts_dir.mkdir(parents=True, exist_ok=True)
        self.seed_bytes = seed.read_bytes()
        self.callback_log = evidence_dir / "callbacks.jsonl"

    def record(self, entry: dict[str, Any]) -> None:
        with self.lock:
            self.seq += 1
            entry["seq"] = self.seq
            self.callbacks.append(entry)
            with self.callback_log.open("a", encoding="utf-8") as fh:
                fh.write(json.dumps(entry, ensure_ascii=False) + "\n")


STATE: ProbeState | None = None


# ---------------------------------------------------------------------------
# HTTP handler
# ---------------------------------------------------------------------------

EDITOR_HTML = """<!doctype html>
<html lang="zh-CN"><head><meta charset="utf-8"><title>OO 9.4 双用户探针 - {user}</title>
<style>html,body,#ph{{height:100%;margin:0}}#bar{{position:fixed;z-index:9;right:8px;top:8px;
background:#fff;border:1px solid #ccc;padding:4px 8px;font:12px/1.6 sans-serif}}</style></head>
<body><div id="bar">user=<b>{user}</b> key=<span id="k"></span> state=<span id="st">boot</span></div>
<div id="ph"></div>
<script src="{oo_url}/web-apps/apps/api/documents/api.js"></script>
<script>
const st = document.getElementById('st');
fetch('/config?user={user}').then(r => r.json()).then(cfg => {{
  document.getElementById('k').textContent = cfg.config.document.key;
  cfg.config.events = {{
    onAppReady: () => st.textContent = 'app-ready',
    onDocumentReady: () => st.textContent = 'doc-ready',
    onError: e => st.textContent = 'error:' + JSON.stringify(e && e.data),
    onWarning: e => st.textContent = 'warning:' + JSON.stringify(e && e.data),
    onDocumentStateChange: e => st.textContent = e.data ? 'dirty' : 'saved',
  }};
  window.__docEditor = new DocsAPI.DocEditor('ph', cfg.config);
}}).catch(e => st.textContent = 'config-failed:' + e);
</script></body></html>
"""


class Handler(BaseHTTPRequestHandler):
    protocol_version = "HTTP/1.1"

    def log_message(self, fmt: str, *args: Any) -> None:  # noqa: A003
        sys.stderr.write("[probe-http] " + (fmt % args) + "\n")

    # -- helpers ----------------------------------------------------------
    def _send(self, code: int, body: bytes, ctype: str) -> None:
        self.send_response(code)
        self.send_header("Content-Type", ctype)
        self.send_header("Content-Length", str(len(body)))
        self.end_headers()
        self.wfile.write(body)

    def _json(self, code: int, obj: Any) -> None:
        self._send(code, json.dumps(obj, ensure_ascii=False).encode("utf-8"), "application/json; charset=utf-8")

    # -- routes -----------------------------------------------------------
    def do_GET(self) -> None:  # noqa: N802
        assert STATE is not None
        parts = urlsplit(self.path)
        route, query = parts.path, parse_qs(parts.query)
        if route == "/health":
            self._json(200, {"ok": True, "doc_key": STATE.doc_key})
        elif route == "/doc":
            self._send(
                200,
                STATE.seed_bytes,
                "application/vnd.openxmlformats-officedocument.spreadsheetml.sheet",
            )
        elif route == "/config":
            user = (query.get("user") or ["alice"])[0]
            self._json(200, build_config(STATE, user))
        elif route == "/editor":
            user = (query.get("user") or ["alice"])[0]
            html = EDITOR_HTML.format(user=user, oo_url=OO_URL_FROM_HOST)
            self._send(200, html.encode("utf-8"), "text/html; charset=utf-8")
        elif route == "/state":
            self._json(200, {"doc_key": STATE.doc_key, "count": len(STATE.callbacks), "callbacks": STATE.callbacks})
        else:
            self._json(404, {"error": "no such probe route", "path": route})

    def do_POST(self) -> None:  # noqa: N802
        assert STATE is not None
        parts = urlsplit(self.path)
        if parts.path == "/control/fail-next":
            count = int((parse_qs(parts.query).get("count") or ["1"])[0])
            with STATE.lock:
                STATE.fail_next = count
            self._json(200, {"fail_next": count})
            return
        if parts.path != "/callback":
            self._json(404, {"error": "no such probe route", "path": parts.path})
            return
        length = int(self.headers.get("Content-Length") or 0)
        raw = self.rfile.read(length) if length else b""
        entry: dict[str, Any] = {
            "ts": _now(),
            "phase": STATE.phase,
            "doc_key_expected": STATE.doc_key,
            "remote_addr": self.client_address[0],
            "query": parse_qs(parts.query),
            "raw_body_sha256": _sha256_bytes(raw),
            "raw_body_len": len(raw),
            "header_names": sorted(self.headers.keys()),
        }
        # --- Authorization header / JWT 实证（Requirement 5.1：claim schema 真校验）---
        auth = self.headers.get("Authorization")
        entry["authorization_present"] = auth is not None
        if auth:
            token = auth[7:] if auth.lower().startswith("bearer ") else auth
            entry["authorization_token_ref"] = _digest_ref(token)
            try:
                claims = jwt.decode(token, OO_SECRET, algorithms=["HS256"])
                entry["jwt_verified"] = True
                entry["jwt_claim_keys"] = sorted(claims.keys())
                entry["jwt_claims"] = redact_payload(claims)
            except Exception as exc:  # noqa: BLE001 — 如实记录，不静默放行
                entry["jwt_verified"] = False
                entry["jwt_error"] = f"{type(exc).__name__}: {exc}"
        try:
            body = json.loads(raw.decode("utf-8")) if raw else {}
        except Exception as exc:  # noqa: BLE001
            entry["body_parse_error"] = f"{type(exc).__name__}: {exc}"
            STATE.record(entry)
            self._json(200, {"error": 1})
            return

        entry["status"] = body.get("status")
        entry["body"] = redact_payload(body)
        entry["body_keys"] = sorted(body.keys())
        entry["has_userdata"] = "userdata" in body
        entry["userdata"] = body.get("userdata")
        entry["users"] = body.get("users")
        entry["actions"] = body.get("actions")

        # --- status 2/6 带 url：下载 artifact 并验证聚合内容 ---
        url = body.get("url")
        if url:
            fetch_url = _rewrite_download_host(url)
            entry["download_url_rewritten_netloc"] = urlsplit(fetch_url).netloc
            try:
                req = urllib.request.Request(fetch_url, method="GET")
                with urllib.request.urlopen(req, timeout=60) as resp:  # noqa: S310 — 固定 OO 主机
                    data = resp.read()
                entry["artifact_bytes"] = len(data)
                entry["artifact_sha256"] = _sha256_bytes(data)
                name = f"{STATE.phase}_cb{STATE.seq + 1:02d}_status{body.get('status')}.xlsx"
                (STATE.artifacts_dir / name).write_bytes(data)
                entry["artifact_file"] = f"artifacts/{name}"
                entry["artifact_probe"] = _xlsx_cell_probe(data, sorted(STATE.marker_cells.values()))
            except Exception as exc:  # noqa: BLE001 — 如实记录下载失败
                entry["artifact_error"] = f"{type(exc).__name__}: {exc}"

        # 故障注入：按需返回 OO 非零 error，实证「host 返回非零 → OO 是否重发同一 status」。
        with STATE.lock:
            inject = STATE.fail_next > 0
            if inject:
                STATE.fail_next -= 1
        entry["response_error_returned"] = 1 if inject else 0
        STATE.record(entry)
        self._json(200, {"error": 1 if inject else 0})


def build_config(state: ProbeState, user: str) -> dict[str, Any]:
    base = f"http://{HOST_FROM_CONTAINER}:{state.port}"
    config = {
        "document": {
            "fileType": "xlsx",
            "key": state.doc_key,
            "title": "task4_probe.xlsx",
            "url": f"{base}/doc?key={state.doc_key}",
            "permissions": {"edit": True, "download": True, "print": True},
        },
        "documentType": "cell",
        "editorConfig": {
            "mode": "edit",
            "lang": "zh-CN",
            "callbackUrl": f"{base}/callback?room=probe&generation=1",
            "user": {"id": user, "name": USERS.get(user, user)},
            "customization": {"forcesave": True, "compactHeader": True},
        },
        "type": "desktop",
    }
    return {"config": {**config, "token": _sign(config)}, "oo_url": OO_URL_FROM_HOST}


# ---------------------------------------------------------------------------
# Command Service
# ---------------------------------------------------------------------------


def call_command_service(payload: dict[str, Any], *, jwt_style: str = "payload") -> dict[str, Any]:
    """调 OO Command Service（`c=forcesave|drop|info|version`）。

    jwt_style:
      - ``payload``：Authorization: Bearer jwt({"payload": body})（DS 文档口径）
      - ``flat``：Authorization: Bearer jwt(body)
      - ``none``：不带 Authorization（用于实证 JWT 是否真被强制）
    """
    url = f"{OO_URL_FROM_HOST}/coauthoring/CommandService.ashx"
    data = json.dumps(payload).encode("utf-8")
    headers = {"Content-Type": "application/json"}
    if jwt_style == "payload":
        headers["Authorization"] = "Bearer " + _sign({"payload": payload})
    elif jwt_style == "flat":
        headers["Authorization"] = "Bearer " + _sign(payload)
    req = urllib.request.Request(url, data=data, headers=headers, method="POST")
    started = time.time()
    try:
        with urllib.request.urlopen(req, timeout=30) as resp:  # noqa: S310 — 固定 OO 主机
            raw = resp.read()
            return {
                "http_status": resp.status,
                "elapsed_ms": round((time.time() - started) * 1000),
                "body": json.loads(raw.decode("utf-8")) if raw else None,
                "jwt_style": jwt_style,
            }
    except urllib.error.HTTPError as exc:
        return {
            "http_status": exc.code,
            "elapsed_ms": round((time.time() - started) * 1000),
            "body": exc.read().decode("utf-8", "replace"),
            "jwt_style": jwt_style,
        }
    except Exception as exc:  # noqa: BLE001
        return {
            "error": f"{type(exc).__name__}: {exc}",
            "elapsed_ms": round((time.time() - started) * 1000),
            "jwt_style": jwt_style,
        }


def oo_build_info() -> dict[str, Any]:
    """OO build 号（healthcheck + /info/info.json，另由 docker exec dpkg -l 交叉验证）。"""
    out: dict[str, Any] = {}
    for name, path in (("healthcheck", "/healthcheck"), ("info", "/info/info.json")):
        try:
            with urllib.request.urlopen(OO_URL_FROM_HOST + path, timeout=10) as resp:  # noqa: S310
                out[name] = resp.read(4000).decode("utf-8", "replace")
        except Exception as exc:  # noqa: BLE001
            out[name] = f"ERROR {type(exc).__name__}: {exc}"
    return out


# ---------------------------------------------------------------------------
# CLI
# ---------------------------------------------------------------------------


def cmd_serve(args: argparse.Namespace) -> int:
    global STATE
    evidence_dir = Path(args.evidence_dir)
    seed = Path(args.seed)
    if not seed.exists():
        print(f"FATAL 种子模板不存在: {seed}")
        return 2
    doc_key = args.doc_key or f"t4probe{int(time.time())}"
    marker_cells = {"alice": args.cell_a, "bob": args.cell_b}
    STATE = ProbeState(
        evidence_dir=evidence_dir,
        seed=seed,
        doc_key=doc_key,
        port=args.port,
        marker_cells=marker_cells,
        phase=args.phase,
    )
    seed_copy = evidence_dir / "seed_document.xlsx"
    shutil.copy2(seed, seed_copy)
    meta = {
        "probe": "task4-oo94-multiuser-callback",
        "phase": args.phase,
        "started_at": _now(),
        "oo_url_from_host": OO_URL_FROM_HOST,
        "host_for_oo": HOST_FROM_CONTAINER,
        "probe_port": args.port,
        "doc_key": doc_key,
        "seed_template_source": str(seed.relative_to(REPO_ROOT)).replace("\\", "/"),
        "seed_sha256": _sha256_bytes(STATE.seed_bytes),
        "seed_bytes": len(STATE.seed_bytes),
        "marker_cells": marker_cells,
        "oo_jwt_enabled": True,
        "oo_secret_ref": _digest_ref(OO_SECRET),
        "oo_build_endpoints": oo_build_info(),
    }
    (evidence_dir / f"run_meta_{args.phase}.json").write_text(
        json.dumps(meta, ensure_ascii=False, indent=2), encoding="utf-8"
    )
    print(json.dumps({"serving": True, "phase": args.phase, "doc_key": doc_key, "port": args.port},
                     ensure_ascii=False))
    httpd = ThreadingHTTPServer(("0.0.0.0", args.port), Handler)
    httpd.serve_forever()
    return 0


def _probe_get(port: int, path: str) -> Any:
    with urllib.request.urlopen(f"http://127.0.0.1:{port}{path}", timeout=15) as resp:  # noqa: S310
        return json.loads(resp.read().decode("utf-8"))


def cmd_command(args: argparse.Namespace) -> int:
    state = _probe_get(args.port, "/state")
    payload: dict[str, Any] = {"c": args.c, "key": state["doc_key"]}
    if args.userdata:
        payload["userdata"] = args.userdata
    if args.users:
        payload["users"] = args.users
    result = call_command_service(payload, jwt_style=args.jwt_style)
    record = {
        "ts": _now(),
        "request": payload,
        "jwt_style": args.jwt_style,
        "result": result,
    }
    log = Path(args.evidence_dir) / "commands.jsonl"
    log.parent.mkdir(parents=True, exist_ok=True)
    with log.open("a", encoding="utf-8") as fh:
        fh.write(json.dumps(record, ensure_ascii=False) + "\n")
    print(json.dumps(record, ensure_ascii=False, indent=2))
    return 0


def cmd_fail_next(args: argparse.Namespace) -> int:
    req = urllib.request.Request(
        f"http://127.0.0.1:{args.port}/control/fail-next?count={args.count}", data=b"", method="POST"
    )
    with urllib.request.urlopen(req, timeout=10) as resp:  # noqa: S310 — 本机探针
        print(resp.read().decode("utf-8"))
    return 0


def cmd_state(args: argparse.Namespace) -> int:
    state = _probe_get(args.port, "/state")
    summary = [
        {
            "seq": c.get("seq"),
            "status": c.get("status"),
            "userdata": c.get("userdata"),
            "users": c.get("users"),
            "actions": c.get("actions"),
            "artifact_sha256": c.get("artifact_sha256"),
            "artifact_cells": (c.get("artifact_probe") or {}).get("cells"),
            "jwt_verified": c.get("jwt_verified"),
            "body_keys": c.get("body_keys"),
        }
        for c in state["callbacks"]
    ]
    print(json.dumps({"doc_key": state["doc_key"], "count": state["count"], "summary": summary},
                     ensure_ascii=False, indent=2))
    return 0


def _content_digest(path: Path) -> dict[str, Any]:
    """artifact 的**内容**指纹（全 sheet 非空单元格排序后 hash）+ 字节指纹 + zip 条目指纹。

    用来区分「内容相同 / 字节不同」——这是判断 status 6 与 status 2 能否按 incoming sha
    折叠成同一 application 的唯一判据。
    """
    import zipfile

    import openpyxl

    data = path.read_bytes()
    out: dict[str, Any] = {"file": path.name, "bytes": len(data), "sha256": _sha256_bytes(data)}
    with zipfile.ZipFile(io.BytesIO(data)) as zf:
        out["zip_entries_sha256"] = _sha256_bytes("\n".join(sorted(zf.namelist())).encode("utf-8"))
        out["zip_entry_count"] = len(zf.namelist())
    wb = openpyxl.load_workbook(io.BytesIO(data), data_only=False)
    cells: list[str] = []
    for ws in wb.worksheets:
        for row in ws.iter_rows():
            for cell in row:
                if cell.value is not None:
                    cells.append(f"{ws.title}!{cell.coordinate}={cell.value!r}")
    wb.close()
    out["nonempty_cell_count"] = len(cells)
    out["content_sha256"] = _sha256_bytes("\n".join(sorted(cells)).encode("utf-8"))
    return out


def cmd_analyze(args: argparse.Namespace) -> int:
    """脱敏既有 callbacks.jsonl + 逐 artifact 做 marker/内容/字节三层指纹比对。"""
    evidence_dir = Path(args.evidence_dir)
    log = evidence_dir / "callbacks.jsonl"
    lines = [json.loads(line) for line in log.read_text(encoding="utf-8").splitlines() if line.strip()]
    for entry in lines:
        if "body" in entry:
            entry["body"] = redact_payload(entry["body"])
        if "jwt_claims" in entry:
            entry["jwt_claims"] = redact_payload(entry["jwt_claims"])
    log.write_text(
        "\n".join(json.dumps(e, ensure_ascii=False) for e in lines) + "\n", encoding="utf-8"
    )

    markers = [f"Z{n}" for n in range(90, 96)]
    artifacts: list[dict[str, Any]] = []
    for entry in lines:
        rel = entry.get("artifact_file")
        if not rel:
            continue
        path = evidence_dir / rel
        if not path.exists():
            continue
        info = _content_digest(path)
        info.update(
            {
                "label": f"{entry.get('phase', 'phase1')}#{entry['seq']}",
                "phase": entry.get("phase", "phase1"),
                "seq": entry["seq"],
                "status": entry.get("status"),
                "userdata": entry.get("userdata"),
                "response_error_returned": entry.get("response_error_returned"),
                "markers": _xlsx_cell_probe(path.read_bytes(), markers).get("cells"),
            }
        )
        artifacts.append(info)
    report = {
        "generated_at": _now(),
        "artifacts": artifacts,
        "content_sha_groups": {},
        "byte_sha_unique_count": len({a["sha256"] for a in artifacts}),
    }
    groups: dict[str, list[str]] = {}
    for art in artifacts:
        groups.setdefault(art["content_sha256"], []).append(art["label"])
    report["content_sha_groups"] = groups
    report["same_content_different_bytes"] = [
        {
            "labels": labels,
            "content_sha256": content_sha,
            "byte_sha256": [a["sha256"] for a in artifacts if a["label"] in labels],
            "statuses": [a["status"] for a in artifacts if a["label"] in labels],
        }
        for content_sha, labels in groups.items()
        if len(labels) > 1 and len({a["sha256"] for a in artifacts if a["label"] in labels}) > 1
    ]
    (evidence_dir / "artifact_analysis.json").write_text(
        json.dumps(report, ensure_ascii=False, indent=2), encoding="utf-8"
    )
    print(json.dumps(report, ensure_ascii=False, indent=2))
    return 0


def cmd_finalize(args: argparse.Namespace) -> int:
    """把收集到的 callback 汇总成可复核的 observations.json（不做任何推断，只做投影）。"""
    evidence_dir = Path(args.evidence_dir)
    lines = [
        json.loads(line)
        for line in (evidence_dir / "callbacks.jsonl").read_text(encoding="utf-8").splitlines()
        if line.strip()
    ]
    commands_path = evidence_dir / "commands.jsonl"
    commands = [
        json.loads(line)
        for line in (commands_path.read_text(encoding="utf-8").splitlines() if commands_path.exists() else [])
        if line.strip()
    ]
    obs = {
        "generated_at": _now(),
        "callback_count": len(lines),
        "statuses_seen": sorted({c.get("status") for c in lines if c.get("status") is not None}),
        "status_with_userdata": sorted(
            {c.get("status") for c in lines if c.get("has_userdata")}
        ),
        "status_without_userdata": sorted(
            {c.get("status") for c in lines if not c.get("has_userdata")}
        ),
        "jwt_verified_all": all(c.get("jwt_verified") for c in lines) if lines else None,
        "body_key_union": sorted({k for c in lines for k in (c.get("body_keys") or [])}),
        "artifact_digests": [
            {
                "seq": c["seq"],
                "status": c.get("status"),
                "userdata": c.get("userdata"),
                "sha256": c.get("artifact_sha256"),
                "cells": (c.get("artifact_probe") or {}).get("cells"),
            }
            for c in lines
            if c.get("artifact_sha256")
        ],
        "commands": [
            {"request": c["request"], "http_status": c["result"].get("http_status"),
             "body": c["result"].get("body"), "jwt_style": c.get("jwt_style")}
            for c in commands
        ],
    }
    (evidence_dir / "observations.json").write_text(
        json.dumps(obs, ensure_ascii=False, indent=2), encoding="utf-8"
    )
    print(json.dumps(obs, ensure_ascii=False, indent=2))
    return 0


def main(argv: list[str] | None = None) -> int:
    ap = argparse.ArgumentParser(description="真实 OO 9.4 两用户 callback 黑盒探针（Task 4）")
    ap.add_argument("--evidence-dir", default=str(DEFAULT_EVIDENCE_DIR))
    ap.add_argument("--port", type=int, default=9991)
    sub = ap.add_subparsers(dest="cmd", required=True)

    p_serve = sub.add_parser("serve", help="启动 document host + callback collector")
    p_serve.add_argument("--seed", default=str(DEFAULT_SEED))
    p_serve.add_argument("--doc-key", default=None)
    p_serve.add_argument("--cell-a", default="Z90")
    p_serve.add_argument("--cell-b", default="Z92")
    p_serve.add_argument("--phase", default="phase1")
    p_serve.set_defaults(func=cmd_serve)

    p_fail = sub.add_parser("fail-next", help="让接下来 N 次 callback 返回 OO 非零 error")
    p_fail.add_argument("--count", type=int, default=1)
    p_fail.set_defaults(func=cmd_fail_next)

    p_cmd = sub.add_parser("command", help="调 OO Command Service")
    p_cmd.add_argument("--c", required=True, choices=["forcesave", "drop", "info", "version", "getForgotten"])
    p_cmd.add_argument("--userdata", default=None)
    p_cmd.add_argument("--users", nargs="*", default=None)
    p_cmd.add_argument("--jwt-style", default="payload", choices=["payload", "flat", "none"])
    p_cmd.set_defaults(func=cmd_command)

    p_state = sub.add_parser("state", help="打印已收集 callback 摘要")
    p_state.set_defaults(func=cmd_state)

    p_an = sub.add_parser("analyze", help="脱敏 callbacks.jsonl + artifact 三层指纹比对")
    p_an.set_defaults(func=cmd_analyze)

    p_fin = sub.add_parser("finalize", help="汇总 observations.json")
    p_fin.set_defaults(func=cmd_finalize)

    args = ap.parse_args(argv)
    return int(args.func(args))


if __name__ == "__main__":
    sys.exit(main())

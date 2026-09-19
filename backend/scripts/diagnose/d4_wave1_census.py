# -*- coding: utf-8 -*-
"""D4 Wave 1 只读核定：T01 列映射 digest + T03 公式分类清册。

用法（仓库根）：
  & .venv/Scripts/python.exe backend/scripts/diagnose/d4_wave1_census.py
"""
from __future__ import annotations

import hashlib
import json
import re
import zipfile
from pathlib import Path

from openpyxl import load_workbook
from openpyxl.utils import get_column_letter

REPO = Path(__file__).resolve().parents[3]
TEMPLATE = REPO / "backend" / "wp_templates" / "D" / "D4 收入底稿.xlsx"
OUT = REPO / ".kiro" / "specs" / "d4-revenue-matrix-bidirectional" / "evidence"
OUT.mkdir(parents=True, exist_ok=True)

BRACKET_RE = re.compile(r"\[\d+\]")


def main() -> None:
    wb = load_workbook(TEMPLATE, data_only=False)
    managed = next(
        name
        for name in wb.sheetnames
        if "D4-2" in name and "主营业务收入明细" in name
    )
    ws = wb[managed]

    header_row = next(
        r
        for r in range(1, 20)
        if ws.cell(r, 1).value == "项目"
        and ws.cell(r, 2).value is not None
        and "1" in str(ws.cell(r, 2).value)
    )
    headers = {get_column_letter(c): ws.cell(header_row, c).value for c in range(1, 23)}
    first_data = header_row + 1
    footer_row = None
    footer_text = None
    for r in range(first_data, ws.max_row + 1):
        v = ws.cell(r, 1).value
        if isinstance(v, str) and "合" in v and "计" in v:
            footer_row = r
            footer_text = v
            break
    assert footer_row is not None
    last_data = footer_row - 1

    sample_formulas: dict[str, str] = {}
    for c in range(1, 23):
        cell = ws.cell(first_data, c)
        if isinstance(cell.value, str) and cell.value.startswith("="):
            sample_formulas[get_column_letter(c)] = cell.value

    field_map: list[dict] = [
        {
            "col": "A",
            "header": headers["A"],
            "column_key": "product",
            "mode": "editable",
            "value_type": "text",
            "json_path": "product",
            "store": True,
            "frontend": "product",
        }
    ]
    for i, col in enumerate(list("BCDEFGHIJKLM")):
        field_map.append(
            {
                "col": col,
                "header": headers[col],
                "column_key": f"month_{i + 1:02d}",
                "mode": "editable",
                "value_type": "amount",
                "json_path": f"months/{i}",
                "store": True,
                "frontend": f"months[{i}]",
            }
        )
    field_map.extend(
        [
            {
                "col": "N",
                "header": headers["N"],
                "column_key": "period_total",
                "mode": "formula",
                "value_type": "amount",
                "json_path": "periodTotal",
                "store": False,
                "frontend": "periodTotal (computed)",
                "formula_sample": sample_formulas.get("N"),
            },
            {
                "col": "O",
                "header": headers["O"],
                "column_key": "audit_adjustment",
                "mode": "editable",
                "value_type": "amount",
                "json_path": "auditAdjustment",
                "store": True,
                "frontend": "auditAdjustment",
            },
            {
                "col": "P",
                "header": headers["P"],
                "column_key": "audited",
                "mode": "formula",
                "value_type": "amount",
                "json_path": "audited",
                "store": False,
                "frontend": "audited (computed)",
                "formula_sample": sample_formulas.get("P"),
            },
            {
                "col": "Q",
                "header": headers["Q"],
                "column_key": "prior_unadjusted",
                "mode": "editable",
                "value_type": "amount",
                "json_path": "priorUnadjusted",
                "store": True,
                "frontend": "priorUnadjusted",
            },
            {
                "col": "R",
                "header": headers["R"],
                "column_key": "prior_adjustment",
                "mode": "editable",
                "value_type": "amount",
                "json_path": "priorAdjustment",
                "store": True,
                "frontend": "priorAdjustment",
            },
            {
                "col": "S",
                "header": headers["S"],
                "column_key": "prior_audited",
                "mode": "formula",
                "value_type": "amount",
                "json_path": "priorAudited",
                "store": False,
                "frontend": "priorAudited (computed)",
                "formula_sample": sample_formulas.get("S"),
            },
            {
                "col": "T",
                "header": headers["T"],
                "column_key": "unadjusted_change_rate",
                "mode": "formula",
                "value_type": "amount",
                "json_path": "unadjustedChangeRate",
                "store": False,
                "frontend": "unadjustedChangeRate (computed)",
                "formula_sample": sample_formulas.get("T"),
            },
            {
                "col": "U",
                "header": headers["U"],
                "column_key": "audited_change_rate",
                "mode": "formula",
                "value_type": "amount",
                "json_path": "auditedChangeRate",
                "store": False,
                "frontend": "auditedChangeRate (computed)",
                "formula_sample": sample_formulas.get("U"),
            },
            {
                "col": "V",
                "header": headers["V"],
                "column_key": "remark",
                "mode": "editable",
                "value_type": "text",
                "json_path": "remark",
                "store": True,
                "frontend": "remark",
            },
        ]
    )
    contract_cols = set(list("ABCDEFGHIJKLM") + ["N", "O", "Q", "R", "V"])
    for fm in field_map:
        fm["contract_candidate"] = fm["col"] in contract_cols

    digest_payload = {
        "contract_fields": [
            {
                "col": x["col"],
                "column_key": x["column_key"],
                "header": x["header"],
                "json_path": x["json_path"],
                "mode": x["mode"],
            }
            for x in field_map
            if x["contract_candidate"]
        ],
        "first_data_row": first_data,
        "footer_marker_exact": footer_text,
        "footer_row": footer_row,
        "header_row": header_row,
        "last_data_row": last_data,
        "managed_sheet": managed,
        "template_relative_path": "D/D4 收入底稿.xlsx",
    }
    canon = json.dumps(digest_payload, ensure_ascii=False, sort_keys=True, separators=(",", ":"))
    mapping_digest = hashlib.sha256(canon.encode("utf-8")).hexdigest()

    mapping = {
        "measured_at": "2026-09-12",
        "tool": "openpyxl data_only=False",
        "template_path": "backend/wp_templates/D/D4 收入底稿.xlsx",
        "template_sha256_preclean": hashlib.sha256(TEMPLATE.read_bytes()).hexdigest(),
        "n_sheets": len(wb.sheetnames),
        "managed_sheet": managed,
        "dimensions": ws.dimensions,
        "header_row": header_row,
        "first_data_row": first_data,
        "last_data_row": last_data,
        "footer_row": footer_row,
        "footer_marker_exact": footer_text,
        "footer_marker_codepoints": [hex(ord(ch)) for ch in (footer_text or "")],
        "headers_A_to_V": headers,
        "sample_formulas_first_data_row": sample_formulas,
        "merged_ranges": [str(m) for m in ws.merged_cells.ranges],
        "field_map_all_22": field_map,
        "contract_field_count": sum(1 for x in field_map if x["contract_candidate"]),
        "contract_fields": [x for x in field_map if x["contract_candidate"]],
        "formula_mask_cols": ["N", "P", "S", "T", "U"],
        "formula_mask_note": (
            "契约 18 条不含 P/S/T/U；它们仍是模板内部公式列，"
            "projection 不得覆盖（需 formula mask / skip）。"
        ),
        "store_item_id": "D4-2-rows",
        "frontend_source": (
            "audit-platform/frontend/src/components/workpaper/"
            "composables/useD4RevenueDetail.ts"
        ),
        "persisted_store_keys": [
            "rowId",
            "product",
            "months",
            "auditAdjustment",
            "priorUnadjusted",
            "priorAdjustment",
            "remark",
        ],
        "mapping_digest": mapping_digest,
        "digest_payload": digest_payload,
    }
    (OUT / "T01-column-field-mapping.json").write_text(
        json.dumps(mapping, ensure_ascii=False, indent=2) + "\n",
        encoding="utf-8",
    )

    # --- T03 formula census ---
    # 权威口径：openpyxl 展开共享公式后的逻辑格（shared master + si 引用都计入）。
    # zip 级另记 master 文本 / 空 shared 引用，供 Task 4 净化脚本定向删 <f>[n]</f>。
    totals = {"external_bracket": 0, "internal": 0, "total_f": 0}
    by_sheet: list[dict] = []
    managed_n: list[dict] = []
    for name in wb.sheetnames:
        ws_i = wb[name]
        ext: list[dict] = []
        inn: list[dict] = []
        for row in ws_i.iter_rows(
            min_row=1, max_row=ws_i.max_row, max_col=ws_i.max_column
        ):
            for cell in row:
                v = cell.value
                if not (isinstance(v, str) and v.startswith("=")):
                    continue
                totals["total_f"] += 1
                item = {"cell": cell.coordinate, "formula": v}
                if BRACKET_RE.search(v):
                    totals["external_bracket"] += 1
                    ext.append(item)
                else:
                    totals["internal"] += 1
                    inn.append(item)
                    if name == managed and cell.column == 14:  # N
                        managed_n.append(item)
        by_sheet.append(
            {
                "sheet": name,
                "external_bracket_count": len(ext),
                "internal_count": len(inn),
                "external_bracket_cells": ext,
                "internal_sample": inn[:30],
                "internal_truncated": len(inn) > 30,
            }
        )

    zip_masters = {
        "external_bracket_f_text": 0,
        "internal_f_text": 0,
        "empty_shared_or_selfclose": 0,
    }
    with zipfile.ZipFile(TEMPLATE) as zf:
        ext_parts = [
            i.filename
            for i in zf.infolist()
            if i.filename.startswith("xl/externalLinks/externalLink")
            and i.filename.endswith(".xml")
        ]
        for info in zf.infolist():
            if not (
                info.filename.startswith("xl/worksheets/sheet")
                and info.filename.endswith(".xml")
            ):
                continue
            xml = zf.read(info.filename).decode("utf-8")
            for m in re.finditer(r"<f([^>]*)(?:/>|>(.*?)</f>)", xml, re.S):
                body = m.group(2)
                if body is None or body == "":
                    zip_masters["empty_shared_or_selfclose"] += 1
                elif BRACKET_RE.search(body):
                    zip_masters["external_bracket_f_text"] += 1
                else:
                    zip_masters["internal_f_text"] += 1

    data_n = [
        x
        for x in managed_n
        if x["cell"].startswith("N")
        and x["cell"][1:].isdigit()
        and 12 <= int(x["cell"][1:]) <= 23
    ]

    formula_census = {
        "measured_at": "2026-09-12",
        "template_path": "backend/wp_templates/D/D4 收入底稿.xlsx",
        "authority": "openpyxl_expanded_shared_formulas",
        "totals": totals,
        "zip_level_f_tag_masters": zip_masters,
        "external_link_parts": ext_parts,
        "external_link_part_count": len(ext_parts),
        "managed_sheet": managed,
        "managed_N_column_internal_formulas": managed_n,
        "managed_N12_N23_sum_formulas": data_n,
        "managed_N_data_rows_all_internal": all(
            not BRACKET_RE.search(x["formula"]) for x in data_n
        ),
        "managed_N_data_rows_are_sum_bm": all(
            x["formula"].startswith("=SUM(B") and ":M" in x["formula"] for x in data_n
        ),
        "by_sheet": by_sheet,
        "classification_rule": (
            "openpyxl 展开后：formula text contains [n] => external_bracket "
            "(sanitize 删 <f> 保 <v>); else internal (含 shared SUM(B:M)，逐字节保留)。"
            "zip 级 shared 空引用继承 master 分类，不得按空 <f/> 误删。"
        ),
    }
    (OUT / "T03-formula-classification-census.json").write_text(
        json.dumps(formula_census, ensure_ascii=False, indent=2) + "\n",
        encoding="utf-8",
    )

    print(
        json.dumps(
            {
                "mapping_digest": mapping_digest,
                "managed_sheet": managed,
                "geometry": {
                    "header_row": header_row,
                    "first_data_row": first_data,
                    "last_data_row": last_data,
                    "footer_row": footer_row,
                    "footer_marker_exact": footer_text,
                },
                "contract_field_count": mapping["contract_field_count"],
                "formula_totals": totals,
                "zip_level_f_tag_masters": zip_masters,
                "external_link_part_count": len(ext_parts),
                "managed_N_count": len(managed_n),
                "managed_N12_N23_count": len(data_n),
                "managed_N_sample": data_n[:3],
                "managed_N_data_rows_all_internal": formula_census[
                    "managed_N_data_rows_all_internal"
                ],
                "managed_N_data_rows_are_sum_bm": formula_census[
                    "managed_N_data_rows_are_sum_bm"
                ],
            },
            ensure_ascii=False,
            indent=2,
        )
    )


if __name__ == "__main__":
    main()

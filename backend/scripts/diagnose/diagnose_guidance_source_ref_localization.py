#!/usr/bin/env python
"""T7 切片 A 第 1 批：证据驱动的 guidance SourceRef 定位器（**只读**）。

回答一个问题：373 份 canonical guidance 文档的每个 section，究竟有多少能挂上
**一条真实可验证**的 SourceRef。不生成、不写入任何 guidance 数据文件。

═══ 为什么这个脚本在结构上无法编造 range ═══

1. ``range`` 永远不是"作者写下的"，而是**某个单元格自身文本**与 section 正文在
   同一套规范化下比对后命中的坐标。命中必须唯一：0 命中 → gap，≥2 命中 → gap。
   "随便挑一个单元格" 在判据上不可能通过。
2. 候选 ref 产出后**一律真跑** ``app.services.guidance_source_refs.validate_source_refs``
   （内部经 G-ID registry）。它会独立重开 workbook、重算文件 sha256、核 sheet 是否
   存在、核 range 是否非空。坐标若是编的，会被 ``range_empty`` /
   ``sheet_not_found`` / ``digest_mismatch`` 打回，只有 ``status=="valid"`` 才计入
   ``attachable_valid``。
3. ``--self-check`` 段落在报告里冻结「正例必过 + 五条反例必被拒」的实测结果。
   没有它，``attachable_valid=0`` 与「我的校验调用坏了」无法区分 —— 这是本脚本
   防假绿的核心装置，缺一条即整份报告 ``harnessVerdict=BROKEN``。

═══ matchPolicy（判据本体，与 evidence 的 matchPolicy 字段逐条对应）═══

规范化 ``skeleton()``：NFKC → 丢弃 Unicode 大类 P/Z/S/C（标点/空白/符号/控制）
→ ASCII 小写。结果只剩字母/数字/CJK，因此全角半角、项目符号、换行、缩进差异
不影响比对，但**不会**因为放宽而把不同正文判成同一段。

分层判据（按序求值，第一个产出候选的层决定结论）：

* T1 ``exact_cell_skeleton``      cell == section
* T2 ``cell_contains_section``    section ⊂ cell，要求 len(section) ≥ 24
* T3 ``section_contains_cell``    cell ⊂ section，要求 len(cell) ≥ 40
* T4 ``longest_common_substring`` 最长公共子串 ≥ 24 且取到该最大值的单元格唯一
* T5 ``multi_token_colocation``   从 section 正文切出的 token（skeleton ≥ 6）中，
  至少 3 个各自在**同一 sheet** 内唯一命中，且只有一个 sheet 达到该法定数；
  range = 这些坐标的外接矩形

sheet 作用域两段式：先 ``matchingSheetNames``（真正的目标 sheet），全层未命中再
退到整册所有 sheet（更宽的搜索面 + 同样的唯一性要求），并记 ``scopeUsed``。

用法：
  python backend/scripts/diagnose/diagnose_guidance_source_ref_localization.py
  python backend/scripts/diagnose/diagnose_guidance_source_ref_localization.py --limit 20
  python backend/scripts/diagnose/diagnose_guidance_source_ref_localization.py --self-check-only
"""
from __future__ import annotations

import argparse
import json
import re
import sys
import unicodedata
from collections import Counter, OrderedDict
from dataclasses import dataclass
from datetime import UTC, datetime
from difflib import SequenceMatcher
from pathlib import Path
from typing import Any, Iterable

ROOT = Path(__file__).resolve().parents[3]
BACKEND = ROOT / "backend"
DIAGNOSE_DIR = Path(__file__).resolve().parent
for _extra in (BACKEND, DIAGNOSE_DIR):
    if str(_extra) not in sys.path:
        sys.path.insert(0, str(_extra))

from openpyxl import load_workbook  # noqa: E402
from openpyxl.utils import get_column_letter  # noqa: E402
from openpyxl.utils.cell import coordinate_to_tuple  # noqa: E402

# 模板 authority 分类**复用**既有可行性脚本，避免第二份映射逻辑分叉。
import diagnose_guidance_source_ref_feasibility as feas  # noqa: E402
from app.services.guidance_gid import GUIDANCE_SECTION_KEYS  # noqa: E402
from app.services.guidance_inventory import (  # noqa: E402
    CANONICAL_SECTION_KEYS,
    GUIDANCE_DIR,
)
from app.services.guidance_source_refs import (  # noqa: E402
    SourceRefContext,
    TemplateAuthority,
    file_sha256,
    validate_source_refs,
)
from app.services.workpaper_sync.canonical_paths import TEMPLATE_ROOT  # noqa: E402

SPEC_DIR = ROOT / ".kiro" / "specs" / "workpaper-guidance-content-closure"
REPORT = SPEC_DIR / "basis" / "T07A-localization-facts.json"
FEASIBILITY_REPORT = SPEC_DIR / "basis" / "T07-source-ref-feasibility.json"

VALIDATOR_ENTRY = "app.services.guidance_source_refs.validate_source_refs"

#: BCD markdown 权威根（memory 登记的路径）。缺失即 bcd_markdown ref 物理不可产。
BCD_ROOT = ROOT / "基础数据" / "致同通用审计程序及底稿模板（2025年修订）" / "BCD类底稿md"

MIN_SECTION_SKELETON = 12
MIN_SECTION_FOR_CONTAINMENT = 24
MIN_CELL_FOR_REVERSE_CONTAINMENT = 40
LCS_MIN = 24
LCS_GRAM = 8
TOKEN_MIN = 6
TOKEN_QUORUM = 3
EXCERPT_MAX = 80

_TOKEN_SPLIT = re.compile(r"[\n\r\u2022•、；;，,。.:：/|—\-()（）\[\]【】“”\"'>《》*#]+")

#: 逐字层：命中即"这段正文的字面就在那一格里"。T4/T5 只是"高度重叠/共位"，
#: 结构上可验证但不能自证 provenance，必须人工确认。
VERBATIM_TIERS = frozenset(
    {"exact_cell_skeleton", "cell_contains_section", "section_contains_cell"}
)


# ═══════════════════════════════════════════════════════════════════════════
# 规范化与判据原语
# ═══════════════════════════════════════════════════════════════════════════


def skeleton(value: str) -> str:
    """把任意正文压成只含字母/数字/CJK 的可比对骨架（matchPolicy 第 1 条）。"""
    text = unicodedata.normalize("NFKC", value or "")
    return "".join(
        ch.lower()
        for ch in text
        if unicodedata.category(ch)[0] not in ("P", "Z", "S", "C")
    )


def excerpt(value: str, limit: int = EXCERPT_MAX) -> str:
    """evidence 里的正文摘录：单行化并硬截断（禁止整段落盘）。"""
    one_line = re.sub(r"\s+", " ", (value or "")).strip()
    return one_line[:limit]


def grams(value: str, size: int = LCS_GRAM) -> set[str]:
    if len(value) < size:
        return set()
    return {value[i : i + size] for i in range(len(value) - size + 1)}


def tokens_of(content: str) -> set[str]:
    """从 section 正文切出可独立定位的 token（matchPolicy T5）。"""
    out: set[str] = set()
    for part in _TOKEN_SPLIT.split(content or ""):
        sk = skeleton(part)
        if len(sk) >= TOKEN_MIN:
            out.add(sk)
    return out


def bounding_range(coordinates: Iterable[str]) -> str:
    """一组坐标的外接矩形（单一 A1 区域，validator 只接受这种形态）。"""
    pairs = [coordinate_to_tuple(c) for c in coordinates]
    rows = [r for r, _ in pairs]
    cols = [c for _, c in pairs]
    return (
        f"{get_column_letter(min(cols))}{min(rows)}:"
        f"{get_column_letter(max(cols))}{max(rows)}"
    )


# ═══════════════════════════════════════════════════════════════════════════
# workbook 单元格快照（read_only，跳过 ~$ 锁文件）
# ═══════════════════════════════════════════════════════════════════════════


@dataclass(frozen=True)
class CellFact:
    sheet: str
    coordinate: str
    text: str
    sk: str


class WorkbookCache:
    """按 path 缓存最近若干册的单元格骨架；顺序处理时命中率高。"""

    def __init__(self, capacity: int = 6) -> None:
        self._capacity = capacity
        self._store: OrderedDict[Path, tuple[list[str], list[CellFact]]] = OrderedDict()
        self.read_errors: dict[str, str] = {}

    def get(self, path: Path) -> tuple[list[str], list[CellFact]]:
        if path in self._store:
            self._store.move_to_end(path)
            return self._store[path]
        value = self._load(path)
        self._store[path] = value
        self._store.move_to_end(path)
        while len(self._store) > self._capacity:
            self._store.popitem(last=False)
        return value

    def _load(self, path: Path) -> tuple[list[str], list[CellFact]]:
        if path.name.startswith("~$"):
            self.read_errors[path.as_posix()] = "lock_file_skipped"
            return [], []
        workbook = None
        sheet_names: list[str] = []
        facts: list[CellFact] = []
        try:
            workbook = load_workbook(path, read_only=True, data_only=False, keep_links=False)
            sheet_names = list(workbook.sheetnames)
            for name in sheet_names:
                for row in workbook[name].iter_rows():
                    for cell in row:
                        raw = cell.value
                        if not isinstance(raw, str) or not raw.strip():
                            continue
                        sk = skeleton(raw)
                        if sk:
                            facts.append(CellFact(name, cell.coordinate, raw, sk))
        except Exception as exc:  # noqa: BLE001 - 读失败必须显式登记，不得静默当空册
            self.read_errors[path.as_posix()] = type(exc).__name__
        finally:
            if workbook is not None:
                workbook.close()
        return sheet_names, facts


# ═══════════════════════════════════════════════════════════════════════════
# 定位分层（唯一命中才产候选；0 或 ≥2 一律 gap）
# ═══════════════════════════════════════════════════════════════════════════


@dataclass(frozen=True)
class Localization:
    tier: str
    sheet: str | None
    cell_range: str | None
    match_count: int
    detail: str
    matched_cells: tuple[str, ...] = ()


def _decide(tier: str, hits: list[CellFact], detail: str) -> Localization:
    if len(hits) == 1:
        hit = hits[0]
        return Localization(
            tier=tier,
            sheet=hit.sheet,
            cell_range=f"{hit.coordinate}:{hit.coordinate}",
            match_count=1,
            detail=detail,
            matched_cells=(f"{hit.sheet}!{hit.coordinate}",),
        )
    return Localization(
        tier=tier,
        sheet=None,
        cell_range=None,
        match_count=len(hits),
        detail=detail,
        matched_cells=tuple(sorted(f"{h.sheet}!{h.coordinate}" for h in hits)[:6]),
    )


def _lcs_size(left: str, right: str) -> int:
    return SequenceMatcher(None, left, right, autojunk=False).find_longest_match(
        0, len(left), 0, len(right)
    ).size


def localize(section_sk: str, content: str, cells: list[CellFact]) -> Localization | None:
    """按 T1→T5 求值；返回 None 表示所有层都 0 命中。"""
    if not cells:
        return None

    t1 = [c for c in cells if c.sk == section_sk]
    if t1:
        return _decide("exact_cell_skeleton", t1, "单元格骨架与 section 正文完全一致")

    if len(section_sk) >= MIN_SECTION_FOR_CONTAINMENT:
        t2 = [c for c in cells if section_sk in c.sk]
        if t2:
            return _decide(
                "cell_contains_section", t2, "单元格文本完整包含 section 正文骨架"
            )

    t3 = [
        c
        for c in cells
        if len(c.sk) >= MIN_CELL_FOR_REVERSE_CONTAINMENT and c.sk in section_sk
    ]
    if t3:
        return _decide(
            "section_contains_cell", t3, "section 正文完整包含该长单元格文本骨架"
        )

    # T4：最长公共子串。LCS ≥ 24 必然共享至少一个 8-gram，故 gram 预筛无损。
    section_grams = grams(section_sk)
    if section_grams:
        best = 0
        winners: list[CellFact] = []
        for cell in cells:
            if len(cell.sk) < LCS_MIN:
                continue
            if not (grams(cell.sk) & section_grams):
                continue
            size = _lcs_size(section_sk, cell.sk)
            if size > best:
                best, winners = size, [cell]
            elif size == best and best > 0:
                winners.append(cell)
        if best >= LCS_MIN:
            return _decide(
                "longest_common_substring",
                winners,
                f"最长公共子串 {best} 字（阈值 {LCS_MIN}）",
            )

    # T5：多 token 同 sheet 共位。token 来自 section 正文本身，且每个必须唯一命中。
    tokens = tokens_of(content)
    if len(tokens) >= TOKEN_QUORUM:
        by_sheet: dict[str, list[CellFact]] = {}
        for cell in cells:
            by_sheet.setdefault(cell.sheet, []).append(cell)
        quorum_sheets: list[tuple[str, list[tuple[str, str]]]] = []
        for sheet, sheet_cells in by_sheet.items():
            resolved: list[tuple[str, str]] = []
            for token in sorted(tokens):
                hits = [c.coordinate for c in sheet_cells if token in c.sk]
                if len(hits) == 1:
                    resolved.append((token, hits[0]))
            if len(resolved) >= TOKEN_QUORUM:
                quorum_sheets.append((sheet, resolved))
        if len(quorum_sheets) == 1:
            sheet, resolved = quorum_sheets[0]
            return Localization(
                tier="multi_token_colocation",
                sheet=sheet,
                cell_range=bounding_range(coord for _, coord in resolved),
                match_count=1,
                detail=(
                    f"{len(resolved)}/{len(tokens)} 个 token 在 sheet 内各自唯一命中"
                    f"（法定数 {TOKEN_QUORUM}）"
                ),
                matched_cells=tuple(f"{sheet}!{coord}" for _, coord in resolved[:6]),
            )
        if quorum_sheets:
            return Localization(
                tier="multi_token_colocation",
                sheet=None,
                cell_range=None,
                match_count=len(quorum_sheets),
                detail="多个 sheet 同时达到 token 法定数，无法唯一归属",
                matched_cells=tuple(sorted(s for s, _ in quorum_sheets)[:6]),
            )
    return None


# ═══════════════════════════════════════════════════════════════════════════
# 模板 authority（复用 feasibility 脚本的分类逻辑，不另写一份）
# ═══════════════════════════════════════════════════════════════════════════


@dataclass(frozen=True)
class Authority:
    kind: str  # single_template | parent_workbook_sheet | ambiguous_template | no_template
    path: Path | None
    matching_sheets: tuple[str, ...] = ()


class AuthorityIndex:
    def __init__(self) -> None:
        paths = feas._template_files()
        self._by_code = feas._index_templates(paths)
        self._sheet_names = {path: feas._sheet_names(path) for path in paths}
        self._sheet_index = feas._index_sheet_codes(paths, self._sheet_names)

    def resolve(self, wp_code: str) -> Authority:
        matches = self._by_code.get(wp_code, [])
        sheet_hits = self._sheet_index.get(wp_code, [])
        if len(matches) == 1:
            template = matches[0]
            own = tuple(name for path, name in sheet_hits if path == template)
            return Authority("single_template", template, own)
        if len(matches) > 1:
            return Authority("ambiguous_template", None)
        distinct = {path for path, _ in sheet_hits}
        if len(distinct) == 1:
            template = next(iter(distinct))
            return Authority(
                "parent_workbook_sheet",
                template,
                tuple(sorted({name for _, name in sheet_hits})),
            )
        if sheet_hits:
            return Authority("ambiguous_template", None)
        return Authority("no_template", None)


def build_context(wp_code: str, template: Path, digest: str) -> SourceRefContext:
    """按冻结模板 lineage 造 validator context；authority 只含这一册。"""
    relative = template.resolve().relative_to(ROOT).as_posix()
    return SourceRefContext(
        target_wp_code=wp_code,
        template_authorities=(
            TemplateAuthority(
                canonical_path=relative,
                active_path=template,
                wp_codes=(wp_code,),
                origin="canonical",
                version=None,
                expected_active_digest=digest,
            ),
        ),
        repo_root=ROOT,
        template_root=TEMPLATE_ROOT,
    )


# ═══════════════════════════════════════════════════════════════════════════
# 校验器自检（反向自检：正例必过、五条反例必被拒）
# ═══════════════════════════════════════════════════════════════════════════


def _run_validator(ref: dict[str, Any], context: SourceRefContext) -> tuple[str, list[str]]:
    batch = validate_source_refs([ref], context)
    result = batch.results[0]
    return result.status, [issue.code for issue in result.issues]


def self_check(index: AuthorityIndex, cache: WorkbookCache) -> dict[str, Any]:
    """在真实模板上冻结「能通过 / 必被拒」实测；缺任何一条 → BROKEN。"""
    probe: tuple[Path, CellFact] | None = None
    for wp_code in sorted(_guidance_codes()):
        authority = index.resolve(wp_code)
        if authority.kind != "single_template" or authority.path is None:
            continue
        _, cells = cache.get(authority.path)
        candidates = [c for c in cells if len(c.sk) >= MIN_SECTION_SKELETON]
        if candidates:
            probe = (authority.path, candidates[0])
            probe_code = wp_code
            break
    if probe is None:
        return {"verdict": "BROKEN", "reason": "no_probe_template_available", "cases": []}

    template, cell = probe
    digest = file_sha256(template)
    relative = template.resolve().relative_to(ROOT).as_posix()
    context = build_context(probe_code, template, digest)
    base = {
        "kind": "xlsx",
        "path": relative,
        "digest": digest,
        "sheet": cell.sheet,
        "range": f"{cell.coordinate}:{cell.coordinate}",
    }

    cases: list[dict[str, Any]] = []

    def record(name: str, ref: dict[str, Any], expect_status: str, expect_code: str | None) -> None:
        status, codes = _run_validator(ref, context)
        ok = status == expect_status and (expect_code is None or expect_code in codes)
        cases.append(
            {
                "case": name,
                "expectStatus": expect_status,
                "expectIssueCode": expect_code,
                "observedStatus": status,
                "observedIssueCodes": codes,
                "pass": ok,
            }
        )

    record("positive_real_cell", dict(base), "valid", None)
    flipped = "0" if digest[-1] != "0" else "1"
    record("negative_digest_flipped", {**base, "digest": digest[:-1] + flipped}, "stale", "digest_mismatch")
    record("negative_sheet_absent", {**base, "sheet": "__no_such_sheet__"}, "stale", "sheet_not_found")
    record("negative_range_empty", {**base, "range": "XFD1048570:XFD1048576"}, "stale", "range_empty")
    record("negative_path_outside_template_root", {**base, "path": "README.md"}, "invalid", "authority_root_not_allowed")
    record(
        "negative_lifecycle_kind",
        {**base, "kind": "custom_candidate"},
        "invalid",
        "kind_is_lifecycle_not_locator",
    )
    verdict = "OK" if all(item["pass"] for item in cases) else "BROKEN"
    return {
        "verdict": verdict,
        "probe": {"wpCode": probe_code, "path": relative, "sheet": cell.sheet, "cell": cell.coordinate},
        "cases": cases,
    }


# ═══════════════════════════════════════════════════════════════════════════
# 主扫描
# ═══════════════════════════════════════════════════════════════════════════


def _guidance_docs() -> list[Path]:
    return [p for p in sorted(GUIDANCE_DIR.glob("*.json")) if not p.stem.startswith("_")]


def _guidance_codes() -> list[str]:
    codes: list[str] = []
    for path in _guidance_docs():
        try:
            document = json.loads(path.read_text(encoding="utf-8"))
        except (OSError, json.JSONDecodeError):
            continue
        if isinstance(document, dict):
            codes.append(str(document.get("wp_code") or path.stem).strip().upper())
    return codes


def _section_scope(authority: Authority, sheet_names: list[str]) -> list[tuple[str, list[str]]]:
    """两段式 sheet 作用域：先目标 sheet，再整册。"""
    scopes: list[tuple[str, list[str]]] = []
    matching = [name for name in authority.matching_sheets if name in sheet_names]
    if matching:
        scopes.append(("matching_sheets", matching))
    if sheet_names and (not matching or len(matching) != len(sheet_names)):
        scopes.append(("all_sheets", list(sheet_names)))
    return scopes


def scan(limit: int | None) -> dict[str, Any]:
    index = AuthorityIndex()
    cache = WorkbookCache()
    checks = self_check(index, cache)

    totals = Counter()
    rejection_codes = Counter()
    tier_hits = Counter()
    by_authority: dict[str, Counter] = {}
    entries: list[dict[str, Any]] = []
    sheet_anomalies: list[dict[str, Any]] = []
    section_key_mismatch = sorted(set(GUIDANCE_SECTION_KEYS) ^ set(CANONICAL_SECTION_KEYS))

    docs = _guidance_docs()
    if limit is not None:
        docs = docs[:limit]

    for doc_path in docs:
        try:
            document = json.loads(doc_path.read_text(encoding="utf-8"))
        except (OSError, json.JSONDecodeError) as exc:
            totals["documents_unreadable"] += 1
            entries.append({"wpCode": doc_path.stem, "error": type(exc).__name__})
            continue
        if not isinstance(document, dict):
            totals["documents_unreadable"] += 1
            continue
        wp_code = str(document.get("wp_code") or doc_path.stem).strip().upper()
        authority = index.resolve(wp_code)
        bucket = by_authority.setdefault(authority.kind, Counter())
        bucket["documents"] += 1
        totals["documents_scanned"] += 1

        sections = [
            item
            for item in (document.get("sections") or [])
            if isinstance(item, dict) and str(item.get("key") or "") in GUIDANCE_SECTION_KEYS
        ]
        entry: dict[str, Any] = {
            "wpCode": wp_code,
            "authority": authority.kind,
            "templatePath": (
                authority.path.resolve().relative_to(ROOT).as_posix() if authority.path else None
            ),
            "matchingSheetNames": list(authority.matching_sheets),
            "canonicalSectionCount": len(sections),
            "sections": [],
        }

        sheet_names: list[str] = []
        cells: list[CellFact] = []
        digest = ""
        context: SourceRefContext | None = None
        if authority.path is not None:
            sheet_names, cells = cache.get(authority.path)
            try:
                digest = file_sha256(authority.path)
            except OSError as exc:
                cache.read_errors[authority.path.as_posix()] = type(exc).__name__
            if digest:
                context = build_context(wp_code, authority.path, digest)
        scopes = _section_scope(authority, sheet_names)

        for section in sections:
            key = str(section.get("key"))
            content = str(section.get("content") or "")
            section_sk = skeleton(content)
            totals["sections_scanned"] += 1
            bucket["sections"] += 1
            row: dict[str, Any] = {
                "sectionKey": key,
                "contentSkeletonLength": len(section_sk),
                "contentExcerpt": excerpt(content),
                "preexistingSourceRefCount": len(section.get("source_refs") or []),
            }

            if context is None or not cells:
                row["verdict"] = "gap_no_template"
                row["gapCode"] = (
                    "ambiguous_template"
                    if authority.kind == "ambiguous_template"
                    else ("no_template" if authority.path is None else "template_unreadable")
                )
                totals["gap_no_template"] += 1
                bucket["gap_no_template"] += 1
                entry["sections"].append(row)
                continue

            if len(section_sk) < MIN_SECTION_SKELETON:
                row["verdict"] = "gap_no_match"
                row["gapCode"] = "content_too_short"
                totals["gap_no_match"] += 1
                bucket["gap_no_match"] += 1
                entry["sections"].append(row)
                continue

            located: Localization | None = None
            scope_used = None
            for scope_name, scope_sheets in scopes:
                subset = [c for c in cells if c.sheet in set(scope_sheets)]
                found = localize(section_sk, content, subset)
                if found is not None:
                    located, scope_used = found, scope_name
                    break
            if located is None:
                row["verdict"] = "gap_no_match"
                row["gapCode"] = "no_unique_cell_located"
                totals["gap_no_match"] += 1
                bucket["gap_no_match"] += 1
                entry["sections"].append(row)
                continue

            row["matchTier"] = located.tier
            row["scopeUsed"] = scope_used
            row["matchDetail"] = located.detail
            row["matchCount"] = located.match_count
            if located.cell_range is None:
                row["verdict"] = "gap_ambiguous"
                row["gapCode"] = f"ambiguous_match:{located.tier}"
                row["ambiguousCells"] = list(located.matched_cells)
                totals["gap_ambiguous"] += 1
                bucket["gap_ambiguous"] += 1
                entry["sections"].append(row)
                continue

            candidate = {
                "kind": "xlsx",
                "path": entry["templatePath"],
                "digest": digest,
                "sheet": located.sheet,
                "range": located.cell_range,
            }
            status, codes = _run_validator(candidate, context)
            row["candidateRef"] = {
                "kind": "xlsx",
                "path": candidate["path"],
                "sheet": candidate["sheet"],
                "range": candidate["range"],
                "digestRecorded": True,
            }
            row["validatorStatus"] = status
            row["validatorIssueCodes"] = codes
            tier_hits[located.tier] += 1
            if status == "valid":
                row["verdict"] = "attachable_valid"
                totals["attachable_valid"] += 1
                bucket["attachable_valid"] += 1
                # validator 只证明「这条 ref 结构上可验证」（sheet 存在 / range 非空 /
                # digest 匹配），**不**证明「这段正文就是从这一格来的」。逐字层
                # (T1–T3) 与启发式层 (T4/T5) 必须分开计数，否则下游会把 18 读成
                # 「18 段已确认 provenance」。
                verbatim = located.tier in VERBATIM_TIERS
                own_sheet = located.sheet in set(authority.matching_sheets)
                row["provenanceStrength"] = "verbatim" if verbatim else "heuristic"
                row["locatedSheetIsOwnCode"] = own_sheet
                row["humanConfirmationRequired"] = (not verbatim) or (not own_sheet)
                totals[
                    "attachable_valid_verbatim" if verbatim else "attachable_valid_heuristic"
                ] += 1
                if not own_sheet:
                    totals["attachable_valid_located_sheet_not_own_code"] += 1
                    sheet_anomalies.append(
                        {
                            "wpCode": wp_code,
                            "sectionKey": key,
                            "templatePath": entry["templatePath"],
                            "ownCodeSheetNames": list(authority.matching_sheets),
                            "locatedSheet": located.sheet,
                            "templateSheetNames": list(sheet_names),
                            "note": "模板内没有承载本 wp_code 的 sheet，候选 ref 只能指向别名/示例 sheet",
                        }
                    )
                if row["humanConfirmationRequired"]:
                    totals["attachable_valid_human_confirmation_required"] += 1
            else:
                row["verdict"] = "validator_rejected"
                totals["validator_rejected"] += 1
                bucket["validator_rejected"] += 1
                for code in codes or ["<no_issue_code>"]:
                    rejection_codes[code] += 1
            entry["sections"].append(row)
        entries.append(entry)

    recorded_counts = {}
    if FEASIBILITY_REPORT.is_file():
        try:
            recorded_counts = json.loads(
                FEASIBILITY_REPORT.read_text(encoding="utf-8")
            ).get("authorityCounts", {})
        except (OSError, json.JSONDecodeError):
            recorded_counts = {}
    derived_counts = {kind: c["documents"] for kind, c in sorted(by_authority.items())}

    return {
        "task": 7,
        "slice": "A1",
        "purpose": "source_ref_localization_facts",
        "recordedAt": datetime.now(UTC).isoformat(),
        "policy": "unique_match_only_no_fabricated_range",
        "readOnly": True,
        "mutatedFiles": [],
        "validatorEntry": VALIDATOR_ENTRY,
        "validatorSelfCheck": checks,
        "harnessVerdict": checks["verdict"],
        "matchPolicy": {
            "normalization": (
                "NFKC → 丢弃 Unicode 大类 P/Z/S/C（标点/空白/符号/控制）→ ASCII 小写；"
                "结果只含字母/数字/CJK"
            ),
            "uniqueness": "命中 0 条 → gap_no_match；命中 ≥2 条 → gap_ambiguous；只有唯一命中才产候选 ref",
            "sheetScope": "先 matchingSheetNames，全层未命中再退整册所有 sheet（scopeUsed 记录实际用的那一段）",
            "tiers": [
                {"tier": "exact_cell_skeleton", "rule": "cell_skeleton == section_skeleton"},
                {
                    "tier": "cell_contains_section",
                    "rule": f"section_skeleton ⊂ cell_skeleton 且 len(section) ≥ {MIN_SECTION_FOR_CONTAINMENT}",
                },
                {
                    "tier": "section_contains_cell",
                    "rule": f"cell_skeleton ⊂ section_skeleton 且 len(cell) ≥ {MIN_CELL_FOR_REVERSE_CONTAINMENT}",
                },
                {
                    "tier": "longest_common_substring",
                    "rule": f"LCS ≥ {LCS_MIN} 且取到最大值的单元格唯一（{LCS_GRAM}-gram 预筛无损）",
                },
                {
                    "tier": "multi_token_colocation",
                    "rule": (
                        f"≥{TOKEN_QUORUM} 个 token（skeleton ≥ {TOKEN_MIN}）在同一 sheet 内各自唯一命中，"
                        "且只有一个 sheet 达到法定数；range = 外接矩形"
                    ),
                },
            ],
            "whyNotFabricable": [
                "range 不是作者写的，而是被比对命中的单元格坐标；唯一性要求让「随便挑一格」不可能通过",
                f"每条候选 ref 都真跑 {VALIDATOR_ENTRY}，它独立重开 workbook、重算 sha256、核 sheet 与非空 range",
                "validatorSelfCheck 的正例证明「若定位成功则能产出 valid」，故 attachable_valid=0 是定位结论而非 harness 故障",
            ],
            "minSectionSkeleton": MIN_SECTION_SKELETON,
        },
        "headline": [
            "attachable_valid 只表示「validator 认这条 ref 结构上可验证」，不表示 provenance 已确认",
            "attachable_valid_verbatim 才是「正文字面就在那一格」；heuristic 层必须人工确认后才可入库",
            "本批不写任何 guidance 数据文件；挂 ref 也不改变 T21 的 pending/complete 计数",
        ],
        "sectionKeyConstant": "app.services.guidance_gid.GUIDANCE_SECTION_KEYS",
        "sectionKeySetMismatchWithInventory": section_key_mismatch,
        "bcdMarkdown": {
            "root": BCD_ROOT.relative_to(ROOT).as_posix(),
            "exists": BCD_ROOT.exists(),
            "markdownFileCount": len(list(BCD_ROOT.rglob("*.md"))) if BCD_ROOT.exists() else 0,
            "note": (
                "bcd_markdown ref 需要 path+digest+headingPath+anchor；源根不在工作树时"
                "任何 headingPath 都无法验证，故本批不产 bcd_markdown 候选"
            ),
        },
        "limitApplied": limit,
        "authorityCrossCheck": {
            "derived": derived_counts,
            "recordedInFeasibilityReport": recorded_counts,
            # 只有全量扫描才可与冻结报告等值比对；--limit 下语义是"部分子集"。
            "comparable": limit is None,
            "consistent": (
                None
                if limit is not None or not recorded_counts
                else derived_counts == recorded_counts
            ),
        },
        "templateSheetNameAnomalies": sheet_anomalies,
        "templateReadErrors": dict(sorted(cache.read_errors.items())),
        "totals": {
            "documents_scanned": totals["documents_scanned"],
            "documents_unreadable": totals["documents_unreadable"],
            "sections_scanned": totals["sections_scanned"],
            "attachable_valid": totals["attachable_valid"],
            "attachable_valid_verbatim": totals["attachable_valid_verbatim"],
            "attachable_valid_heuristic": totals["attachable_valid_heuristic"],
            "attachable_valid_located_sheet_not_own_code": totals[
                "attachable_valid_located_sheet_not_own_code"
            ],
            "attachable_valid_human_confirmation_required": totals[
                "attachable_valid_human_confirmation_required"
            ],
            "gap_no_match": totals["gap_no_match"],
            "gap_ambiguous": totals["gap_ambiguous"],
            "gap_no_template": totals["gap_no_template"],
            "validator_rejected": totals["validator_rejected"],
            "validator_rejection_codes": dict(sorted(rejection_codes.items())),
            "candidate_tier_hits": dict(sorted(tier_hits.items())),
        },
        "byAuthority": {
            kind: {name: value for name, value in sorted(counter.items())}
            for kind, counter in sorted(by_authority.items())
        },
        "entries": entries,
    }


def main() -> int:
    for stream in (sys.stdout, sys.stderr):
        try:
            stream.reconfigure(encoding="utf-8", errors="replace")  # type: ignore[union-attr]
        except (AttributeError, ValueError):
            pass

    parser = argparse.ArgumentParser(description="guidance SourceRef 定位事实核算（只读）")
    parser.add_argument("--limit", type=int, default=None, help="只扫前 N 份 guidance 文档")
    parser.add_argument("--self-check-only", action="store_true", help="只跑校验器自检")
    parser.add_argument("--out", type=Path, default=REPORT, help="evidence JSON 输出路径")
    args = parser.parse_args()

    if args.self_check_only:
        checks = self_check(AuthorityIndex(), WorkbookCache())
        print(json.dumps(checks, ensure_ascii=False, indent=2))
        return 0 if checks["verdict"] == "OK" else 1

    report = scan(args.limit)
    out: Path = args.out
    out.parent.mkdir(parents=True, exist_ok=True)
    out.write_text(json.dumps(report, ensure_ascii=False, indent=2) + "\n", encoding="utf-8")

    totals = report["totals"]
    print(f"harnessVerdict={report['harnessVerdict']}")
    print(
        "docs={documents_scanned} sections={sections_scanned} "
        "attachable_valid={attachable_valid} gap_no_match={gap_no_match} "
        "gap_ambiguous={gap_ambiguous} gap_no_template={gap_no_template} "
        "validator_rejected={validator_rejected}".format(**totals)
    )
    print(
        "attachable_valid split: verbatim={attachable_valid_verbatim} "
        "heuristic={attachable_valid_heuristic} "
        "sheet_not_own_code={attachable_valid_located_sheet_not_own_code} "
        "needs_human={attachable_valid_human_confirmation_required}".format(**totals)
    )
    print(f"rejection_codes={totals['validator_rejection_codes']}")
    print(f"tier_hits={totals['candidate_tier_hits']}")
    print(f"authorityCrossCheck={report['authorityCrossCheck']}")
    print(f"report={out}")
    return 0 if report["harnessVerdict"] == "OK" else 2


if __name__ == "__main__":
    raise SystemExit(main())

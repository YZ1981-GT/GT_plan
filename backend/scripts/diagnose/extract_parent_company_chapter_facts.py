"""母公司附注章事实基线提取器（只读）。

判据真源 = ``docs/模版/`` 下两份源 docx，**不是** ``note_template_*.json``（后者是 md 重建产物）。

用法::

    python backend/scripts/diagnose/extract_parent_company_chapter_facts.py
    python backend/scripts/diagnose/extract_parent_company_chapter_facts.py --out backend/data/parent_note_source_facts.json

设计要点（踩坑已登记，改动前先读）:

- **docx 章号是 Word 自动编号，段落文本不含「十二、」** → 定位母公司章必须按 ``paragraph.style.name``
  是 ``Heading 1`` / ``标题 1``；用 ``^\\d+、`` 正则匹配章号会 0 命中。
- 章范围 = 该 Heading 1 到**下一个 Heading 1**之间的 body 元素区间；表格靠 body 顺序归属到最近的
  上游 Heading，故必须按 ``doc.element.body`` 顺序遍历，**不能**分别遍历 ``doc.paragraphs`` 与
  ``doc.tables``（那样丢失相对位置）。
- **两版章的子节 Heading 层级不同**：listed 母公司章子节是 H2、合并章子节是 H2；soe 母公司章子节是
  H3、合并章子节是 H3，且两版都有更深一层（H4/H5）的子子节。故「子节」= 该章内**最浅**的标题层级，
  更深层标题的表必须**上卷**到所属子节。soe 长期股权投资的 3 张表就分布在 H3 + 两个 H4 下
  （5 列 / 7 列 / 12 列），不上卷会数成「1 张表」而与需求 4.4 的 3 张表判据冲突。
- 「同构子节」有**两级依据**：listed 写在**子节标题**里（「披露格式参考附注五、X」），soe 写在
  **章首说明**里（「应参照上述相应项目的要求加以注释」）。故章首说明必须单独承载
  （``intro_paragraphs`` / ``chapter_reference_text``），不能塞进 ``sections`` 当哨兵 ——
  那样会让 ``len(sections)`` 两版语义不同，且 soe 的参照声明在第 2 段、只对第 1 段求值会判成 False。
- 判定「是否同构」的**结构性**判据是上卷后 ``own_table_count == 0``，参照声明只作旁证。
- 合并章体量大（listed 82 子节 / soe 93 子节），只对「母公司对应科目」输出完整表结构，
  其余子节只输出标题 + 行列数摘要。
- 控制台在 Windows 是 GBK，stdout 只打 ASCII 摘要，中文明细一律走 ``--out`` 落 UTF-8 JSON。
- 源 docx 缺失时**明确报错并打印期望路径**，不得静默返回空。

_Requirements: 1.1, 1.2, 1.4_
"""

from __future__ import annotations

import argparse
import json
import re
import sys
from dataclasses import asdict, dataclass, field
from pathlib import Path
from typing import Any

# ── 源 docx 路径（单一真源，勿散落在别处） ──────────────────────────────────
REPO_ROOT = Path(__file__).resolve().parents[3]

LISTED_DOCX = (
    REPO_ROOT
    / "docs"
    / "模版"
    / "1.上市公司年审报表及附注-2026.01"
    / "1.上市公司年审报表及附注-2026.01"
    / "3.2025年度上市公司财务报表附注模板-2026.01.15.docx"
)
SOE_DOCX = (
    REPO_ROOT
    / "docs"
    / "模版"
    / "1、2025年度财务决算审计报告-2026.01.06"
    / "1、2025年度财务决算审计报告-国企"
    / "1.1-2025国企财务报表附注20260119.docx"
)

SOURCE_DOCX: dict[str, Path] = {"listed": LISTED_DOCX, "soe": SOE_DOCX}

# 母公司章标题（源 docx 原文）。
# listed 侧源 docx 不含「母」字，而 note_template_listed.json 写的是「母公司财务报表主要项目注释」
# —— 该偏差已裁决保留 JSON 现值（requirements 需求 3），此处只如实记录源 docx 原值。
PARENT_CHAPTER_TITLE: dict[str, str] = {
    "listed": "公司财务报表主要项目注释",
    "soe": "母公司财务报表的主要项目附注",
}

# 合并章标题（供「同构子节」比对）
CONSOLIDATED_CHAPTER_TITLE: dict[str, str] = {
    "listed": "合并财务报表项目附注",
    "soe": "财务报表主要项目注释",
}

# 母公司子节 → 合并章子节的标题别名。
# 归一（去空白 + 去括注）后 listed 4 个同构子节与合并章**逐字相等**，无需别名；
# soe 只有一处用语差异：母公司章写「营业收入与营业成本」、合并章写「营业收入、营业成本」。
COUNTERPART_TITLE_ALIASES: dict[str, dict[str, str]] = {
    "listed": {},
    "soe": {"营业收入与营业成本": "营业收入、营业成本"},
}

# 「披露格式参考附注五、X」这类同构声明
_REF_PATTERNS = (
    re.compile(r"披露格式参考附注[^）)]*"),
    re.compile(r"参照(?:上述)?相应项目[^；;。)）]*"),
)

_HEADING_RE = re.compile(r"^(?:Heading|标题)\s*(\d+)$")
_PAREN_RE = re.compile(r"[（(][^（()）]*[)）]")


class SourceDocxMissingError(FileNotFoundError):
    """源 docx 不存在。判据失效必须显式报错，不得静默返回空。"""


class ChapterNotFoundError(LookupError):
    """在源 docx 中找不到指定 Heading 1 章。"""


# ── 数据模型 ────────────────────────────────────────────────────────────────
@dataclass
class TableFact:
    """源 docx 里一张表的结构事实。"""

    index_in_section: int
    rows: int
    cols: int
    two_level_header: bool
    header_row1: list[str] = field(default_factory=list)
    # 只在 two_level_header 为真时给值；单级表给 None（避免把首个数据行误当第二行表头）
    header_row2: list[str] | None = None
    first_col_sample: list[str] = field(default_factory=list)
    # 该表所属的最深层标题（soe 长期股权投资的 7/12 列子表挂在 H4「对子公司投资」等之下）
    owner_heading: str | None = None
    owner_heading_level: int | None = None


@dataclass
class SectionFact:
    """一个子节（该章最浅标题层级）的事实，表已按更深层标题上卷。"""

    heading_level: int
    title: str
    normalized_title: str
    body_index: int
    # 参照声明（listed 在标题里，soe 在章首说明里 → 后者见 ChapterFact.chapter_reference_text）
    has_reference_note: bool
    reference_text: str | None
    # 上卷后的自有表数量；== 0 即「同构子节」（结构性判据）
    own_table_count: int
    is_isomorphic: bool
    sub_headings: list[str] = field(default_factory=list)
    tables: list[TableFact] = field(default_factory=list)
    paragraph_samples: list[str] = field(default_factory=list)


@dataclass
class ChapterFact:
    variant: str
    chapter_title: str
    body_index: int
    body_span: tuple[int, int]
    # 该章子节所在的标题层级（listed 母公司章 = 2，soe 母公司章 = 3）
    top_child_level: int | None
    heading1_titles: list[str] = field(default_factory=list)
    # 章首说明（第一个子节标题之前的正文段落）。soe 的「应参照上述相应项目的要求加以注释」在此，
    # 是需求 1.4「同构子节」的**章级**依据，必须与子节分开承载。
    intro_paragraphs: list[str] = field(default_factory=list)
    chapter_reference_text: str | None = None
    sections: list[SectionFact] = field(default_factory=list)


# ── docx 遍历 ───────────────────────────────────────────────────────────────
def _heading_level(style_name: str | None) -> int | None:
    """把样式名翻成 Heading 层级；非标题返回 None。

    同时认英文 ``Heading N`` 与中文 ``标题 N``（同一份 docx 内可能混用）。
    """
    if not style_name:
        return None
    m = _HEADING_RE.match(str(style_name).strip())
    return int(m.group(1)) if m else None


def _iter_body(doc: Any) -> list[tuple[str, Any]]:
    """按 body 顺序返回 ``[(kind, obj)]``，kind ∈ {'p','tbl'}。"""
    from docx.table import Table
    from docx.text.paragraph import Paragraph

    out: list[tuple[str, Any]] = []
    for child in doc.element.body.iterchildren():
        tag = child.tag.rsplit("}", 1)[-1]
        if tag == "p":
            out.append(("p", Paragraph(child, doc)))
        elif tag == "tbl":
            out.append(("tbl", Table(child, doc)))
    return out


def _norm_space(text: str | None) -> str:
    return re.sub(r"\s+", " ", (text or "")).strip()


def _norm_title(text: str | None) -> str:
    """标题归一：去全部空白 + 去括注。

    listed 母公司章标题带「（披露格式参考附注五、4）」、合并章标题带「（以下不适用的，请删除…）」，
    去括注后两侧才可比。
    """
    t = re.sub(r"\s+", "", text or "")
    prev = None
    while prev != t:  # 嵌套括注需反复剥
        prev = t
        t = _PAREN_RE.sub("", t)
    return t


def _cell_text(cell: Any) -> str:
    return _norm_space(cell.text)


def _row_tcs(row: Any) -> list[Any]:
    """返回该行各 cell 的底层 ``<w:tc>`` 元素列表。

    🔴 必须把元素**存进列表持有引用**再比较 ``id()``。写成
    ``[id(c._tc) for c in row.cells]`` 时 ``_Cell`` 与 lxml 代理是临时对象，
    比较前可能已被回收、id 被复用 → 合并单元格判定随机失真。
    """
    return [c._tc for c in row.cells]


def _table_fact(tbl: Any, idx: int, *, owner: str | None, owner_level: int | None) -> TableFact:
    rows = list(tbl.rows)
    n_rows = len(rows)
    n_cols = len(tbl.columns) if n_rows else 0

    h1 = [_cell_text(c) for c in rows[0].cells] if n_rows >= 1 else []

    # 两级表头判据：第 1 行存在横向合并（相邻 cell 指向同一 tc），或第 1 行与第 2 行同列
    # 指向同一 tc（rowspan=2 的独立列，两级表头里的「项目」「合计」就是这种）。
    two_level = False
    if n_rows >= 2 and n_cols >= 2:
        row0 = _row_tcs(rows[0])
        row1 = _row_tcs(rows[1])
        has_hmerge = len({id(t) for t in row0}) < len(row0)
        has_vmerge = any(a is b for a, b in zip(row0, row1))
        two_level = bool(has_hmerge or has_vmerge)

    h2 = [_cell_text(c) for c in rows[1].cells] if (two_level and n_rows >= 2) else None

    first_col = [r.cells[0].text and _cell_text(r.cells[0]) or "" for r in rows[: min(n_rows, 40)] if r.cells]

    return TableFact(
        index_in_section=idx,
        rows=n_rows,
        cols=n_cols,
        two_level_header=two_level,
        header_row1=h1,
        header_row2=h2,
        first_col_sample=first_col,
        owner_heading=owner,
        owner_heading_level=owner_level,
    )


def _reference_of(text: str) -> str | None:
    for pat in _REF_PATTERNS:
        m = pat.search(text)
        if m:
            return m.group(0)
    return None


def _find_heading1_positions(items: list[tuple[str, Any]]) -> list[tuple[int, str]]:
    out: list[tuple[int, str]] = []
    for i, (kind, obj) in enumerate(items):
        if kind != "p":
            continue
        if _heading_level(getattr(getattr(obj, "style", None), "name", None)) == 1:
            txt = re.sub(r"\s+", "", obj.text or "")
            if txt:
                out.append((i, txt))
    return out


def extract_chapter(doc: Any, *, variant: str, chapter_title: str) -> ChapterFact:
    """从 docx 中抽出指定 Heading 1 章的完整结构事实（表按最浅子节层级上卷）。"""
    items = _iter_body(doc)
    h1_positions = _find_heading1_positions(items)
    h1_titles = [t for _i, t in h1_positions]

    want = re.sub(r"\s+", "", chapter_title)
    target = next((p for p in h1_positions if p[1] == want), None)
    if target is None:  # 二级判据：允许包含匹配（源 docx 标题可能带括注）
        target = next((p for p in h1_positions if want in p[1] or p[1] in want), None)
    if target is None:
        raise ChapterNotFoundError(
            f"{variant}: 源 docx 中未找到 Heading 1 「{chapter_title}」。"
            f" 实际 Heading 1 清单（{len(h1_positions)} 个）: "
            + " | ".join(f"idx={i} {t}" for i, t in h1_positions)
        )

    start = target[0]
    end = next((i for i, _t in h1_positions if i > start), len(items))

    # 先确定「子节层级」= 该章内最浅的标题层级（listed 母公司章 = 2，soe 母公司章 = 3）
    child_levels = [
        lvl
        for kind, obj in items[start + 1 : end]
        if kind == "p"
        and (lvl := _heading_level(getattr(getattr(obj, "style", None), "name", None))) is not None
        and lvl > 1
        and _norm_space(obj.text)
    ]
    top_level = min(child_levels) if child_levels else None

    chapter = ChapterFact(
        variant=variant,
        chapter_title=target[1],
        body_index=start,
        body_span=(start, end),
        top_child_level=top_level,
        heading1_titles=h1_titles,
    )

    cur: SectionFact | None = None
    cur_owner: str | None = None
    cur_owner_level: int | None = None
    tbl_seq = 0

    for i in range(start + 1, end):
        kind, obj = items[i]
        if kind == "p":
            lvl = _heading_level(getattr(getattr(obj, "style", None), "name", None))
            text = _norm_space(obj.text)
            if not text:
                continue
            if lvl is not None and lvl > 1 and top_level is not None:
                if lvl == top_level:
                    ref = _reference_of(text)
                    cur = SectionFact(
                        heading_level=lvl,
                        title=text,
                        normalized_title=_norm_title(text),
                        body_index=i,
                        has_reference_note=ref is not None,
                        reference_text=ref,
                        own_table_count=0,
                        is_isomorphic=False,  # 遍历结束后按 own_table_count 定稿
                    )
                    chapter.sections.append(cur)
                    cur_owner, cur_owner_level, tbl_seq = None, None, 0
                else:  # 更深层标题：记名并把其后的表上卷到当前子节
                    cur_owner, cur_owner_level = text, lvl
                    if cur is not None:
                        cur.sub_headings.append(text)
                continue
            # 正文段落
            if cur is None:
                chapter.intro_paragraphs.append(text)
                if chapter.chapter_reference_text is None:
                    chapter.chapter_reference_text = _reference_of(text)
            elif len(cur.paragraph_samples) < 8:
                cur.paragraph_samples.append(text)
        else:  # tbl
            if cur is None:
                continue
            cur.tables.append(
                _table_fact(obj, tbl_seq, owner=cur_owner, owner_level=cur_owner_level)
            )
            cur.own_table_count += 1
            tbl_seq += 1

    for s in chapter.sections:
        s.is_isomorphic = s.own_table_count == 0

    return chapter


# ── 合并章对应科目（供「同构子节」比对） ────────────────────────────────────
def _counterpart_key(variant: str, parent_section_title: str) -> str:
    """母公司子节标题 → 合并章标题的比对键（归一 + 别名）。"""
    norm = _norm_title(parent_section_title)
    return COUNTERPART_TITLE_ALIASES.get(variant, {}).get(norm, norm)


def summarize_consolidated(
    consol: ChapterFact, *, variant: str, wanted: list[str]
) -> dict[str, Any]:
    """合并章摘要：对母公司对应科目输出完整表结构，其余只输出标题 + 行列数。

    ``wanted`` 是母公司子节标题（原文）；比对走归一 + 别名，避免括注差异导致 0 命中。
    """
    want_keys = {_counterpart_key(variant, t) for t in wanted}
    counterparts: dict[str, Any] = {}
    others: list[dict[str, Any]] = []

    for s in consol.sections:
        entry = {
            "heading_level": s.heading_level,
            "title": s.title,
            "normalized_title": s.normalized_title,
            "own_table_count": s.own_table_count,
            "table_cols": [t.cols for t in s.tables],
            "table_rows": [t.rows for t in s.tables],
        }
        if s.normalized_title in want_keys:
            counterparts.setdefault(
                s.normalized_title,
                {**entry, "sub_headings": s.sub_headings, "tables": [asdict(t) for t in s.tables]},
            )
        else:
            others.append(entry)

    return {
        "chapter_title": consol.chapter_title,
        "body_index": consol.body_index,
        "top_child_level": consol.top_child_level,
        "section_count": len(consol.sections),
        # 母公司同构子节在合并章的对应科目（完整表结构，需求 4.1 的比对依据）
        "counterparts": counterparts,
        "counterparts_missing": sorted(want_keys - set(counterparts)),
        "other_sections": others,
    }


# ── 主流程 ──────────────────────────────────────────────────────────────────
def _load_doc(path: Path, *, variant: str) -> Any:
    if not path.exists():
        raise SourceDocxMissingError(
            f"{variant} 源 docx 不存在。期望路径: {path}"
            " —— 判据真源必须是 docs/模版/ 下的源 docx，不得回退到 note_template_*.json。"
        )
    try:
        from docx import Document
    except ImportError as exc:  # pragma: no cover
        raise SystemExit(f"[ERR] 缺少 python-docx: {exc}") from exc
    return Document(str(path))


def build_facts() -> dict[str, Any]:
    """提取两版母公司章 + 合并章对应科目的完整事实基线。"""
    out: dict[str, Any] = {"source_docx": {}, "parent_chapter": {}, "consolidated_chapter": {}}
    for variant, path in SOURCE_DOCX.items():
        doc = _load_doc(path, variant=variant)
        out["source_docx"][variant] = {
            "path": str(path.relative_to(REPO_ROOT)).replace("\\", "/"),
            "size": path.stat().st_size,
            "parent_chapter_title": PARENT_CHAPTER_TITLE[variant],
            "consolidated_chapter_title": CONSOLIDATED_CHAPTER_TITLE[variant],
        }
        parent = extract_chapter(doc, variant=variant, chapter_title=PARENT_CHAPTER_TITLE[variant])
        out["parent_chapter"][variant] = asdict(parent)

        consol = extract_chapter(
            doc, variant=variant, chapter_title=CONSOLIDATED_CHAPTER_TITLE[variant]
        )
        out["consolidated_chapter"][variant] = summarize_consolidated(
            consol,
            variant=variant,
            wanted=[s.title for s in parent.sections if s.is_isomorphic],
        )
    return out


def _ascii_summary(facts: dict[str, Any]) -> str:
    """ASCII-only 摘要（Windows 控制台是 GBK，中文会 UnicodeEncodeError）。"""
    lines: list[str] = []
    for variant in ("listed", "soe"):
        ch = facts["parent_chapter"].get(variant) or {}
        secs = ch.get("sections", [])
        iso = [s for s in secs if s["is_isomorphic"]]
        own = [s for s in secs if not s["is_isomorphic"]]
        lines.append(
            f"[{variant}] parent chapter: body_index={ch.get('body_index')} "
            f"span={tuple(ch.get('body_span') or ())} child_level=H{ch.get('top_child_level')} "
            f"heading1_count={len(ch.get('heading1_titles') or [])}"
        )
        lines.append(
            f"  sections={len(secs)} isomorphic={len(iso)} own_table={len(own)} "
            f"intro_paras={len(ch.get('intro_paragraphs') or [])} "
            f"chapter_ref={'Y' if ch.get('chapter_reference_text') else 'N'}"
        )
        for n, s in enumerate(secs, 1):
            kind = "ISOMORPHIC" if s["is_isomorphic"] else "OWN-TABLE "
            cols = ",".join(str(t["cols"]) for t in s["tables"]) or "-"
            rows = ",".join(str(t["rows"]) for t in s["tables"]) or "-"
            two = "".join("T" if t["two_level_header"] else "f" for t in s["tables"]) or "-"
            lines.append(
                f"    #{n} {kind} H{s['heading_level']} tables={s['own_table_count']} "
                f"cols=[{cols}] rows=[{rows}] two_level=[{two}] "
                f"title_ref={'Y' if s['has_reference_note'] else 'N'} "
                f"sub_headings={len(s['sub_headings'])}"
            )
        cc = facts["consolidated_chapter"].get(variant) or {}
        lines.append(
            f"  consolidated: child_level=H{cc.get('top_child_level')} "
            f"sections={cc.get('section_count')} "
            f"counterparts={len(cc.get('counterparts') or {})} "
            f"missing={len(cc.get('counterparts_missing') or [])}"
        )
        for key, cp in (cc.get("counterparts") or {}).items():
            cols = ",".join(str(c) for c in cp["table_cols"]) or "-"
            lines.append(
                f"    counterpart#{len(key)}c H{cp['heading_level']} "
                f"tables={cp['own_table_count']} cols=[{cols}]"
            )
        if cc.get("counterparts_missing"):
            lines.append("    [WARN] some isomorphic sections have no counterpart in consolidated")
    return "\n".join(lines)


def main(argv: list[str] | None = None) -> int:
    ap = argparse.ArgumentParser(description="母公司附注章事实基线提取（只读）")
    ap.add_argument("--out", help="把完整事实落成 UTF-8 JSON（中文明细只在此文件里）")
    args = ap.parse_args(argv)

    try:
        facts = build_facts()
    except (SourceDocxMissingError, ChapterNotFoundError) as exc:
        # 判据失效必须显式失败（Requirements 1.2），不得静默返回空
        print(f"[ERR] {exc}", file=sys.stderr)
        return 2

    if args.out:
        p = Path(args.out)
        p.parent.mkdir(parents=True, exist_ok=True)
        p.write_text(json.dumps(facts, ensure_ascii=False, indent=2), encoding="utf-8")
        print(f"[OK] facts written: {p}")

    print(_ascii_summary(facts))
    return 0


if __name__ == "__main__":
    sys.exit(main())

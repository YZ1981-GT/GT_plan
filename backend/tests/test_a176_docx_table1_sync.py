"""Tests for A17-6 docx Table[1] agenda blob roundtrip."""

from __future__ import annotations

import tempfile
from pathlib import Path

import pytest

from app.services.a176_docx_sync import (
    AGENDA_TITLES,
    TEMPLATE_PATH,
    _format_agenda_blob,
    _parse_agenda_blob,
    parse_docx,
)


def _build_minimal_a176_docx(path: Path, agenda_blob: str) -> None:
    """Construct minimal 11-row table doc mimicking A17-6 Table[1]."""
    from docx import Document

    doc = Document()
    doc.add_paragraph("提示横幅")
    table = doc.add_table(rows=11, cols=3)
    table.rows[0].cells[0].text = "被审计单位名称：测试公司"
    table.rows[0].cells[2].text = "索引号：A17-6"
    table.rows[1].cells[0].text = "报表截止日/会计期间：2025年度"
    table.rows[4].cells[0].text = "会议地点：会议室A"
    table.rows[10].cells[0].text = agenda_blob
    doc.save(str(path))


class TestAgendaBlobRoundtrip:
    def test_format_parse_roundtrip(self):
        agenda = {
            1: "审计过程总体顺利",
            5: "发现 2 项未更正错报",
            9: "拟发表无保留意见",
        }
        blob = _format_agenda_blob(agenda)
        parsed = _parse_agenda_blob(blob)
        assert parsed[1] == "审计过程总体顺利"
        assert parsed[5] == "发现 2 项未更正错报"
        assert parsed[9] == "拟发表无保留意见"

    def test_empty_agenda_blob(self):
        blob = _format_agenda_blob({})
        parsed = _parse_agenda_blob(blob)
        assert all(parsed.get(i, "") == "" for i in range(1, 11))
        for i in range(1, 11):
            assert f"{i}、" in blob
            assert AGENDA_TITLES[i] in blob


class TestParseDocxTable1:
    def test_roundtrip_via_tempfile_minimal_doc(self):
        agenda_in = {2: "收入确认风险已应对", 7: "独立性声明已全部签署"}
        blob = _format_agenda_blob(agenda_in)

        with tempfile.TemporaryDirectory() as tmp:
            doc_path = Path(tmp) / "A17-6.docx"
            _build_minimal_a176_docx(doc_path, blob)

            parsed_agenda, meta = parse_docx(doc_path)
            assert parsed_agenda[2] == "收入确认风险已应对"
            assert parsed_agenda[7] == "独立性声明已全部签署"
            assert "测试公司" in meta.get("client_name", "")

            re_blob = _format_agenda_blob(parsed_agenda)
            re_parsed = _parse_agenda_blob(re_blob)
            assert re_parsed[2] == agenda_in[2]
            assert re_parsed[7] == agenda_in[7]

    @pytest.mark.skipif(not TEMPLATE_PATH.exists(), reason="A17-6 template not present")
    def test_roundtrip_with_real_template(self):
        agenda_in = {1: "总结会预填测试", 10: "无其他事项"}
        blob = _format_agenda_blob(agenda_in)

        with tempfile.TemporaryDirectory() as tmp:
            doc_path = Path(tmp) / "A17-6.docx"
            import shutil

            shutil.copy2(TEMPLATE_PATH, doc_path)

            from docx import Document

            doc = Document(str(doc_path))
            table = doc.tables[1] if len(doc.tables) >= 2 else doc.tables[0]
            if len(table.rows) > 10:
                table.rows[10].cells[0].text = blob
            doc.save(str(doc_path))

            parsed_agenda, _meta = parse_docx(doc_path)
            assert parsed_agenda.get(1) == "总结会预填测试"

            re_parsed = _parse_agenda_blob(_format_agenda_blob(parsed_agenda))
            assert re_parsed[1] == agenda_in[1]

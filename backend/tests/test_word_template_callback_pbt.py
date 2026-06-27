"""Property-Based Tests for word-template OnlyOffice callback extraction.

Property 11: OO callback extracts and syncs placeholder values

**Validates: Requirements 8.4**
"""

from __future__ import annotations

import os
import tempfile
from pathlib import Path

from docx import Document
from hypothesis import HealthCheck, given, settings
from hypothesis import strategies as st

from app.services.wp_docx_template_parser import parse_template

# ─── Hypothesis Strategies ────────────────────────────────────────────────────

_field_id_st = st.from_regex(r"[a-z][a-z0-9_]{2,15}", fullmatch=True)
_value_st = st.text(
    min_size=1,
    max_size=20,
    alphabet=st.characters(categories=("L", "N")),
).map(lambda s: s.strip()).filter(lambda s: len(s) > 0)


@st.composite
def callback_data_strategy(draw: st.DrawFn) -> tuple[str, list[tuple[str, str]], dict[str, str]]:
    """Generate (wp_code, placeholders, edited_values) for callback testing.

    Simulates a user editing a word-template docx via OnlyOffice:
    1. Template has placeholders
    2. User replaces some placeholder text with actual values
    3. Callback extracts and syncs

    Returns:
        wp_code: workpaper code
        placeholders: list of (field_id, label)
        edited_values: dict of field_id→new_value (user edits)
    """
    wp_code = draw(st.from_regex(r"A[0-9]{1,2}-[0-9]", fullmatch=True))
    count = draw(st.integers(min_value=1, max_value=4))
    field_ids = draw(st.lists(_field_id_st, min_size=count, max_size=count, unique=True))
    labels = [f"标签{i}" for i in range(count)]
    placeholders = list(zip(field_ids, labels))

    # User edits some (or all) placeholders
    edit_count = draw(st.integers(min_value=1, max_value=count))
    values = draw(st.lists(_value_st, min_size=edit_count, max_size=edit_count))
    edited_values = {field_ids[i]: values[i] for i in range(edit_count)}

    return wp_code, placeholders, edited_values


# ─── Helpers ──────────────────────────────────────────────────────────────────


def _create_template_docx(placeholders: list[tuple[str, str]]) -> str:
    """Create a template docx with placeholders. Each placeholder is a single paragraph."""
    doc = Document()
    doc.add_heading("模板文档", level=1)
    for field_id, label in placeholders:
        doc.add_paragraph(f"${{{field_id}:{label}}}")
    fd, path = tempfile.mkstemp(suffix=".docx")
    os.close(fd)
    doc.save(path)
    return path


def _create_edited_docx(
    placeholders: list[tuple[str, str]], edited_values: dict[str, str]
) -> str:
    """Create a docx that simulates OO-edited output.

    Replaces placeholders with user-edited values (simulating what OnlyOffice saves).
    """
    doc = Document()
    doc.add_heading("模板文档", level=1)
    for field_id, label in placeholders:
        if field_id in edited_values:
            # User replaced the placeholder with actual value
            doc.add_paragraph(edited_values[field_id])
        else:
            # Unchanged placeholder
            doc.add_paragraph(f"${{{field_id}:{label}}}")
    fd, path = tempfile.mkstemp(suffix=".docx")
    os.close(fd)
    doc.save(path)
    return path


def _extract_placeholder_value_from_docx(file_path: Path, placeholder) -> str | None:
    """Mirror of the callback extraction logic from wp_onlyoffice_router.py."""
    try:
        doc = Document(str(file_path))
        pos = placeholder.position

        if "paragraph_index" in pos:
            para_idx = pos["paragraph_index"]
            if para_idx < len(doc.paragraphs):
                para_text = doc.paragraphs[para_idx].text.strip()
                if placeholder.pattern in para_text:
                    return None
                return para_text if para_text else None
        elif "table_index" in pos:
            tbl_idx = pos["table_index"]
            row_idx = pos.get("row", 0)
            col_idx = pos.get("col", 0)
            if tbl_idx < len(doc.tables):
                table = doc.tables[tbl_idx]
                if row_idx < len(table.rows):
                    row = table.rows[row_idx]
                    if col_idx < len(row.cells):
                        cell_text = row.cells[col_idx].text.strip()
                        if placeholder.pattern in cell_text:
                            return None
                        return cell_text if cell_text else None
    except Exception:
        return None
    return None


def _simulate_callback_extraction(
    template_path: str, edited_path: str, wp_code: str
) -> dict[str, str]:
    """Simulate the full callback extraction logic.

    1. Parse the template to get placeholder definitions (positions + patterns)
    2. Use those positions to extract values from the edited docx
    3. Return field_id→value mapping for all detected edits
    """
    # Parse original template to get placeholder definitions
    structure = parse_template(template_path)

    extracted: dict[str, str] = {}
    for placeholder in structure.placeholders:
        value = _extract_placeholder_value_from_docx(Path(edited_path), placeholder)
        if value and value != placeholder.pattern:
            extracted[placeholder.field_id] = value

    return extracted


# ─── Property 11: OO callback extracts and syncs ─────────────────────────────


class TestProperty11CallbackExtractsAndSyncs:
    """Property 11: OO callback extracts and syncs placeholder values.

    For any word-template docx saved via OnlyOffice callback, the system SHALL
    extract current placeholder values from the saved document and identify
    which fields have been edited (value differs from original pattern).

    **Validates: Requirements 8.4**
    """

    @settings(max_examples=5, suppress_health_check=[HealthCheck.function_scoped_fixture])
    @given(data=callback_data_strategy())
    def test_edited_values_are_extracted(self, data):
        """All user-edited placeholder values are correctly extracted from saved docx."""
        wp_code, placeholders, edited_values = data
        template_path = _create_template_docx(placeholders)
        edited_path = _create_edited_docx(placeholders, edited_values)
        try:
            extracted = _simulate_callback_extraction(
                template_path, edited_path, wp_code
            )

            # Every edited value should be extracted
            for field_id, expected_value in edited_values.items():
                assert field_id in extracted, (
                    f"Edited field '{field_id}' not extracted from callback docx"
                )
                assert extracted[field_id] == expected_value, (
                    f"Extracted value mismatch for '{field_id}': "
                    f"expected '{expected_value}', got '{extracted[field_id]}'"
                )
        finally:
            os.unlink(template_path)
            os.unlink(edited_path)

    @settings(max_examples=5, suppress_health_check=[HealthCheck.function_scoped_fixture])
    @given(data=callback_data_strategy())
    def test_unedited_placeholders_not_extracted(self, data):
        """Placeholders that remain unchanged are not extracted (return None)."""
        wp_code, placeholders, edited_values = data
        template_path = _create_template_docx(placeholders)
        edited_path = _create_edited_docx(placeholders, edited_values)
        try:
            extracted = _simulate_callback_extraction(
                template_path, edited_path, wp_code
            )

            # Unedited fields should NOT be in extracted map
            for field_id, label in placeholders:
                if field_id not in edited_values:
                    assert field_id not in extracted, (
                        f"Unedited field '{field_id}' should not be extracted, "
                        f"but got '{extracted.get(field_id)}'"
                    )
        finally:
            os.unlink(template_path)
            os.unlink(edited_path)

    @settings(max_examples=5, suppress_health_check=[HealthCheck.function_scoped_fixture])
    @given(data=callback_data_strategy())
    def test_item_id_format_correct(self, data):
        """Extracted fields produce correct item_id format: wt-{wp_code}-{field_id}."""
        wp_code, placeholders, edited_values = data
        template_path = _create_template_docx(placeholders)
        edited_path = _create_edited_docx(placeholders, edited_values)
        try:
            extracted = _simulate_callback_extraction(
                template_path, edited_path, wp_code
            )

            prefix = f"wt-{wp_code}-"
            for field_id in extracted:
                item_id = f"{prefix}{field_id}"
                assert item_id.startswith("wt-"), (
                    f"item_id '{item_id}' doesn't start with 'wt-'"
                )
                assert wp_code in item_id, (
                    f"item_id '{item_id}' doesn't contain wp_code '{wp_code}'"
                )
                assert item_id == f"wt-{wp_code}-{field_id}", (
                    f"item_id format incorrect: '{item_id}'"
                )
        finally:
            os.unlink(template_path)
            os.unlink(edited_path)

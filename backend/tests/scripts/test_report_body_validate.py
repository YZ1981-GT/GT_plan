"""报告正文 POC 校验冒烟 + 畸形占位符判据。"""

from __future__ import annotations

import re
from pathlib import Path

import pytest
from docx import Document

from scripts.validate_report_body_template import (
    MALFORMED_TOKEN_RE,
    REQUIRED_TOKENS,
    iter_text_locations,
    validate,
)

_BACKEND = Path(__file__).resolve().parent.parent.parent
REPORT_BODY_DIR = _BACKEND / "data/audit_report_templates/report_body"
POC = (
    REPORT_BODY_DIR
    / "1.1 模板A-无保留意见审计报告模板（上市公司、三板创新层及公开发债）-简版.docx"
)


def test_report_body_poc_passes():
    issues = validate(POC)
    assert issues == [], issues


def test_poc_firm_name_placeholder_is_well_formed():
    """回归：POC 的 `{{firm_name}}` 必须是完整双花括号。

    🔴 2026-09-28 这条长期红的根因是**模板数据缺陷**（非判据陈旧）：
    para[21] 实为 `{{firm_name}`，少一个 `}`。原校验只报含糊的
    `missing required token`，分不清"写错了"还是"压根没写"。
    """
    doc = Document(POC)
    full = "\n".join(text for _, text in iter_text_locations(doc))
    assert "{{firm_name}}" in full
    # 反向：不得残留"只有单个右括号"的畸形形态
    assert full.count("{{firm_name}") == full.count("{{firm_name}}")


@pytest.mark.parametrize(
    ("sample", "should_flag"),
    [
        ("{{firm_name}}", False),
        ("{{firm_name}", True),
        ("{firm_name}}", True),
        ("{{a}} 与 {{b}}", False),
        ("普通正文，无占位符", False),
        ("##OPT:key_audit_matters:", False),
    ],
)
def test_malformed_token_detector_both_directions(sample, should_flag):
    """双向变异：畸形形态必须命中，合法形态必须不命中。

    只做单向（"畸形能命中"）会把正常模板判成缺陷；只做另一向则检不出真缺陷。
    """
    assert bool(MALFORMED_TOKEN_RE.search(sample)) is should_flag


def test_all_report_body_templates_have_no_malformed_placeholders():
    """全目录收口：34 本模板（现算）一本都不得带畸形占位符。

    只校 POC 一本会漏掉 `standalone/` 下的同名副本 —— 本次实测两本同缺陷。
    """
    files = sorted(REPORT_BODY_DIR.rglob("*.docx"))
    assert files, f"未扫到任何模板：{REPORT_BODY_DIR}"

    offenders: list[str] = []
    for path in files:
        doc = Document(path)
        for loc, text in iter_text_locations(doc):
            for match in MALFORMED_TOKEN_RE.finditer(text):
                rel = path.relative_to(REPORT_BODY_DIR).as_posix()
                offenders.append(f"{rel} {loc} {match.group(0)!r}")
    assert offenders == [], offenders


def test_required_tokens_are_all_well_formed_patterns():
    """必需 token 清单自身不得写成畸形形态（否则守卫会要求一个错的东西）。"""
    for token in REQUIRED_TOKENS:
        if token.startswith("{{"):
            assert re.fullmatch(r"\{\{[a-z_][\w]*\}\}", token), token
            assert not MALFORMED_TOKEN_RE.search(token), token


def test_validator_reads_table_cells_not_only_paragraphs():
    """`doc.paragraphs` 不含表格内段落 ⇒ 扫描口径必须覆盖单元格。

    变异证明：真源 POC 的表格内现算 0 个占位符，所以用"位置标签里有 table[" 来
    证明扫描器**确实走过**表格通路，而不是靠真源恰好为空蒙过去。
    """
    doc = Document(POC)
    locations = iter_text_locations(doc)
    assert any(loc.startswith("table[") for loc, _ in locations), (
        "未产出任何表格位置 —— 表格通路没走到"
    )
    assert any(loc.startswith("para[") for loc, _ in locations)

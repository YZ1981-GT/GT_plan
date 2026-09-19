"""Tests for AdjudicationExportTemplateService._build_multi_row_header.

Validates: Requirements 3.1, 3.2, 3.3, 3.4, 3.5 (multi-row header structure)
"""

from __future__ import annotations

from openpyxl import Workbook
from openpyxl.styles import Alignment, Border, Font, Side

from app.services.adjudication_export_template_service import (
    AdjudicationExportTemplateService,
)


def _make_ws():
    """Create a fresh worksheet for testing."""
    wb = Workbook()
    return wb.active


class TestBuildMultiRowHeaderNoAging:
    """Test header generation without aging columns (standard 12-column layout)."""

    def test_row2_level1_headers(self):
        """Sub-task 2.1: Row 2 contains correct level-1 header values."""
        ws = _make_ws()
        AdjudicationExportTemplateService._build_multi_row_header(ws, start_col=1)

        assert ws.cell(row=2, column=1).value == "项目"
        assert ws.cell(row=2, column=2).value == "期初"
        assert ws.cell(row=2, column=6).value == "期末"
        assert ws.cell(row=2, column=10).value == "变动额"
        assert ws.cell(row=2, column=11).value == "变动率"
        assert ws.cell(row=2, column=12).value == "原因分析"

    def test_row3_level2_headers(self):
        """Sub-task 2.2: Row 3 contains correct level-2 sub-headers."""
        ws = _make_ws()
        AdjudicationExportTemplateService._build_multi_row_header(ws, start_col=1)

        # 期初 sub-headers (cols 2-5)
        assert ws.cell(row=3, column=2).value == "未审"
        assert ws.cell(row=3, column=3).value == "AJE"
        assert ws.cell(row=3, column=4).value == "RJE"
        assert ws.cell(row=3, column=5).value == "审定"

        # 期末 sub-headers (cols 6-9)
        assert ws.cell(row=3, column=6).value == "未审"
        assert ws.cell(row=3, column=7).value == "AJE"
        assert ws.cell(row=3, column=8).value == "RJE"
        assert ws.cell(row=3, column=9).value == "审定"

    def test_horizontal_merges(self):
        """Sub-task 2.3: 期初 and 期末 are horizontally merged across 4 cols."""
        ws = _make_ws()
        AdjudicationExportTemplateService._build_multi_row_header(ws, start_col=1)

        merged = [str(m) for m in ws.merged_cells.ranges]
        # 期初: row 2, cols 2-5
        assert "B2:E2" in merged
        # 期末: row 2, cols 6-9
        assert "F2:I2" in merged

    def test_vertical_merges(self):
        """Sub-task 2.3: 项目/变动额/变动率/原因分析 are vertically merged."""
        ws = _make_ws()
        AdjudicationExportTemplateService._build_multi_row_header(ws, start_col=1)

        merged = [str(m) for m in ws.merged_cells.ranges]
        # 项目: rows 2-3, col 1
        assert "A2:A3" in merged
        # 变动额: rows 2-3, col 10
        assert "J2:J3" in merged
        # 变动率: rows 2-3, col 11
        assert "K2:K3" in merged
        # 原因分析: rows 2-3, col 12
        assert "L2:L3" in merged

    def test_total_column_count(self):
        """Property 9: Non-aging table has exactly 12 columns."""
        ws = _make_ws()
        AdjudicationExportTemplateService._build_multi_row_header(ws, start_col=1)

        # Last column with a value in row 2 should be 12 (原因分析)
        assert ws.cell(row=2, column=12).value == "原因分析"
        assert ws.cell(row=2, column=13).value is None

    def test_styles_applied(self):
        """Sub-task 2.4: All header cells have bold font, center alignment, thin border."""
        ws = _make_ws()
        AdjudicationExportTemplateService._build_multi_row_header(ws, start_col=1)

        for row in (2, 3):
            for col in range(1, 13):
                cell = ws.cell(row=row, column=col)
                assert cell.font.bold is True, f"Cell ({row},{col}) not bold"
                assert cell.alignment.horizontal == "center", (
                    f"Cell ({row},{col}) not center-aligned"
                )
                assert cell.alignment.vertical == "center", (
                    f"Cell ({row},{col}) not vertically centered"
                )
                assert cell.border.left.style == "thin", (
                    f"Cell ({row},{col}) missing left border"
                )
                assert cell.border.right.style == "thin", (
                    f"Cell ({row},{col}) missing right border"
                )
                assert cell.border.top.style == "thin", (
                    f"Cell ({row},{col}) missing top border"
                )
                assert cell.border.bottom.style == "thin", (
                    f"Cell ({row},{col}) missing bottom border"
                )


class TestBuildMultiRowHeaderWithAging:
    """Test header generation with aging columns."""

    def test_aging_headers_inserted(self):
        """Aging headers appear after col 9 (after 期末审定)."""
        ws = _make_ws()
        aging = ["1年以内", "1-2年", "2-3年", "3年以上"]
        AdjudicationExportTemplateService._build_multi_row_header(
            ws, start_col=1, aging_headers=aging
        )

        # Aging headers start at col 10
        assert ws.cell(row=2, column=10).value == "1年以内"
        assert ws.cell(row=2, column=11).value == "1-2年"
        assert ws.cell(row=2, column=12).value == "2-3年"
        assert ws.cell(row=2, column=13).value == "3年以上"

    def test_tail_columns_shifted(self):
        """变动额/变动率/原因分析 shift after aging columns."""
        ws = _make_ws()
        aging = ["1年以内", "1-2年", "2-3年", "3年以上"]
        AdjudicationExportTemplateService._build_multi_row_header(
            ws, start_col=1, aging_headers=aging
        )

        # tail_start = 10 + 4 = 14
        assert ws.cell(row=2, column=14).value == "变动额"
        assert ws.cell(row=2, column=15).value == "变动率"
        assert ws.cell(row=2, column=16).value == "原因分析"

    def test_total_column_count_with_aging(self):
        """Property 9: Aging table has 12 + len(aging_headers) columns."""
        ws = _make_ws()
        aging = ["1年以内", "1-2年", "2-3年"]
        AdjudicationExportTemplateService._build_multi_row_header(
            ws, start_col=1, aging_headers=aging
        )

        # Total = 12 + 3 = 15
        total_cols = 12 + len(aging)
        assert ws.cell(row=2, column=total_cols).value == "原因分析"
        assert ws.cell(row=2, column=total_cols + 1).value is None

    def test_aging_columns_vertically_merged(self):
        """Aging columns are vertically merged across rows 2-3."""
        ws = _make_ws()
        aging = ["1年以内", "1-2年"]
        AdjudicationExportTemplateService._build_multi_row_header(
            ws, start_col=1, aging_headers=aging
        )

        merged = [str(m) for m in ws.merged_cells.ranges]
        # Col 10 (J) and Col 11 (K) should be vertically merged
        assert "J2:J3" in merged
        assert "K2:K3" in merged

    def test_vertical_merges_shifted_with_aging(self):
        """变动额/变动率/原因分析 vertical merges shift with aging columns."""
        ws = _make_ws()
        aging = ["1年以内", "1-2年"]
        AdjudicationExportTemplateService._build_multi_row_header(
            ws, start_col=1, aging_headers=aging
        )

        merged = [str(m) for m in ws.merged_cells.ranges]
        # tail_start = 10 + 2 = 12
        # 变动额 = col 12 (L), 变动率 = col 13 (M), 原因分析 = col 14 (N)
        assert "L2:L3" in merged  # 变动额
        assert "M2:M3" in merged  # 变动率
        assert "N2:N3" in merged  # 原因分析

    def test_styles_applied_with_aging(self):
        """All cells including aging columns get proper styling."""
        ws = _make_ws()
        aging = ["1年以内", "1-2年"]
        AdjudicationExportTemplateService._build_multi_row_header(
            ws, start_col=1, aging_headers=aging
        )

        total_cols = 12 + len(aging)  # 14
        for row in (2, 3):
            for col in range(1, total_cols + 1):
                cell = ws.cell(row=row, column=col)
                assert cell.font.bold is True, f"Cell ({row},{col}) not bold"
                assert cell.border.left.style == "thin", (
                    f"Cell ({row},{col}) missing border"
                )

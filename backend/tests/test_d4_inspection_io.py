"""D4-13/15/16 导入导出 IO 守卫 — Property 1

验证：
- D4-15 CompletenessItem 三维嵌套的 parse ↔ export 往返一致
- D4-16 英文 key 行结构的 parse ↔ export 往返一致，派生列重算
- D4-13 双 item 文本锚点在 _SPECIAL_ITEM_IDS 中映射正确
- D4-15 item_id = "D4-15-items"
- _SHEET_HEADERS 与专用 parser 字段覆盖一致

更新：Task 3 重写 — 旧 _export_d4_15_row/_export_d4_16_row (singular)
已被 _export_d4_15_rows/_export_d4_16_rows (plural, 批量) 替代，
parser 字段从 productName/quantity 改为 product/qty（三层嵌套英文 key）。
"""
from __future__ import annotations

import io
from typing import Any

import pytest
from openpyxl import Workbook, load_workbook

from app.routers.wp_render_strategies._d4_import_export import (
    _SHEET_HEADERS,
    _SPECIAL_ITEM_IDS,
    _d4_15_is_consistent,
    _d4_16_recompute_derived,
    _export_d4_15_rows,
    _export_d4_16_rows,
    _parse_d4_15_row,
    _parse_d4_16_row,
    _safe_float,
    _resolve_item_id,
)


def _wb_bytes(wb: Workbook) -> bytes:
    buf = io.BytesIO()
    wb.save(buf)
    return buf.getvalue()


def _read_back(data: bytes):
    wb = load_workbook(io.BytesIO(data), read_only=True, data_only=True)
    ws = wb.active
    headers = [str(c.value).strip() if c.value else "" for c in next(ws.iter_rows(min_row=1, max_row=1))]
    rows = list(ws.iter_rows(min_row=2, values_only=True))
    wb.close()
    return headers, rows


class TestD4_13_ItemIds:
    """D4-13 双 item 文本锚点。"""

    def test_process_item_id(self):
        assert _SPECIAL_ITEM_IDS.get("D4-13-process") == "D4-13-process"

    def test_conclusion_item_id(self):
        assert _SPECIAL_ITEM_IDS.get("D4-13-conclusion") == "D4-13-conclusion"


class TestD4_15_IO:
    """D4-15 完整性检查表 IO 往返。"""

    def test_item_id_mapping(self):
        assert _resolve_item_id("D4-15") == "D4-15-items"

    def test_headers_count(self):
        assert len(_SHEET_HEADERS["D4-15"]) == 18

    def test_parse_roundtrip(self):
        """export → parse 往返后字段一致。"""
        item = {
            "id": "c-test",
            "delivery": {"date": "2025-01-15", "number": "FH-001", "product": "产品A", "qty": 100.0, "amount": 50000.0},
            "invoice": {"date": "2025-01-16", "number": "FP-001", "product": "产品A", "qty": 100.0, "amount": 50000.0},
            "voucher": {"date": "2025-01-17", "number": "PZ-001", "product": "产品A", "qty": 100.0, "amount": 50000.0},
            "isConsistent": True,
            "remark": "无异常",
        }
        wb = Workbook()
        ws = wb.active
        ws.append(_SHEET_HEADERS["D4-15"])
        _export_d4_15_rows(ws, [item])
        headers, rows = _read_back(_wb_bytes(wb))
        assert len(rows) == 1
        parsed = _parse_d4_15_row(rows[0], headers)
        assert parsed is not None
        assert parsed["delivery"]["date"] == "2025-01-15"
        assert parsed["delivery"]["product"] == "产品A"
        assert parsed["delivery"]["amount"] == 50000.0
        assert parsed["invoice"]["number"] == "FP-001"
        assert parsed["voucher"]["amount"] == 50000.0
        assert parsed["isConsistent"] is True
        assert parsed["remark"] == "无异常"

    def test_empty_row_skip(self):
        """全空行返回 None。"""
        headers = _SHEET_HEADERS["D4-15"]
        row = tuple([None] * len(headers))
        assert _parse_d4_15_row(row, headers) is None

    def test_dynamic_id_generated(self):
        """导入时 id 由 parser 分配。"""
        headers = _SHEET_HEADERS["D4-15"]
        row = (None,) + ("2025-01-01", "FH-1", "品A", "10", 100.0) + (None,) * 12
        parsed = _parse_d4_15_row(row, headers)
        assert parsed is not None
        assert parsed["id"].startswith("c-")


class TestD4_16_IO:
    """D4-16 出口口岸核对 IO 往返。"""

    def test_item_id_mapping(self):
        assert _resolve_item_id("D4-16") == "D4-16-rows"

    def test_headers_count(self):
        assert len(_SHEET_HEADERS["D4-16"]) == 10

    def test_parse_roundtrip(self):
        """export → parse 往返后字段一致，派生列重算。"""
        item = {
            "id": "e-test",
            "bookAmount": 1000000.0,
            "portsPeriod": "2025-01~06",
            "portsAmount": 950000.0,
            "portsDiff": -50000.0,
            "portsReason": "汇率差异",
            "taxReportAmount": 980000.0,
            "taxDiff": -20000.0,
            "taxReason": "时间性差异",
            "taxIndex": "D4-16-1",
        }
        wb = Workbook()
        ws = wb.active
        ws.append(_SHEET_HEADERS["D4-16"])
        _export_d4_16_rows(ws, [item])
        headers, rows = _read_back(_wb_bytes(wb))
        assert len(rows) == 1
        parsed = _parse_d4_16_row(rows[0], headers)
        assert parsed is not None
        assert parsed["bookAmount"] == 1000000.0
        assert parsed["portsAmount"] == 950000.0
        # 派生列重算
        assert parsed["portsDiff"] == 950000.0 - 1000000.0  # -50000
        assert parsed["taxDiff"] == 980000.0 - 1000000.0    # -20000
        assert parsed["portsReason"] == "汇率差异"
        assert parsed["taxIndex"] == "D4-16-1"

    def test_diff_recomputed(self):
        """重算函数不信任输入的 diff 值。"""
        row: dict[str, Any] = {
            "bookAmount": 500000,
            "portsAmount": 480000,
            "taxReportAmount": 490000,
            "portsDiff": 99999,  # 故意错
            "taxDiff": 88888,    # 故意错
        }
        _d4_16_recompute_derived(row)
        assert row["portsDiff"] == -20000.0  # 480000 - 500000
        assert row["taxDiff"] == -10000.0    # 490000 - 500000

    def test_empty_row_skip(self):
        headers = _SHEET_HEADERS["D4-16"]
        row = tuple([None] * len(headers))
        assert _parse_d4_16_row(row, headers) is None


class TestReverseCheck:
    """反向自检：篡改后必须检测到不一致。"""

    def test_d4_15_consistency_false_on_mismatch(self):
        """D4-15 金额不等时 isConsistent 必须为 False。"""
        item = {
            "delivery": {"amount": 100},
            "invoice": {"amount": 99},
            "voucher": {"amount": 100},
        }
        assert _d4_15_is_consistent(item) is False

    def test_d4_16_diff_changes_on_book_change(self):
        """D4-16 改 bookAmount → diff 重算结果不同。"""
        row: dict[str, Any] = {"bookAmount": 1000, "portsAmount": 800, "taxReportAmount": 900}
        _d4_16_recompute_derived(row)
        original_ports_diff = row["portsDiff"]
        # 篡改 bookAmount
        row["bookAmount"] = 2000
        _d4_16_recompute_derived(row)
        assert row["portsDiff"] != original_ports_diff

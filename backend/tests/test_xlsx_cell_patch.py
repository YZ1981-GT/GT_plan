"""只改目标格的 xlsx 补丁器（spec post-push-residual-defect-closure R4）。

判据：改写落到目标格、样式保留；其余部件逐字节不变（含 x14 数据验证、宏）；共享公式不改；
公式改成文字时删计算链；含 & / < 的文字写成合法 XML；找不到页签时原文件不动。
"""
from __future__ import annotations

import hashlib
import io
import re
import shutil
import zipfile
from pathlib import Path

import openpyxl
import pytest
from openpyxl.styles import Font

from app.services import xlsx_cell_patch as xp
from app.services.xlsx_cell_patch import CellPatch, XlsxPatchError

TEMPLATES = Path(__file__).resolve().parents[1] / "wp_templates"


def _parts(path: Path) -> dict[str, bytes]:
    with zipfile.ZipFile(path) as zf:
        return {n: zf.read(n) for n in zf.namelist()}


def _sheet_part(path: Path, title: str) -> str:
    from app.services.workpaper_sync.excel_extract import _sheet_parts

    with zipfile.ZipFile(path) as zf:
        return _sheet_parts(zf)[title]


def _rewrite_part(path: Path, name: str, transform) -> None:
    entries = _parts(path)
    entries[name] = transform(entries[name].decode("utf-8")).encode("utf-8")
    buf = io.BytesIO()
    with zipfile.ZipFile(buf, "w", zipfile.ZIP_DEFLATED) as out:
        for n, payload in entries.items():
            out.writestr(n, payload)
    path.write_bytes(buf.getvalue())


def _add_part(path: Path, name: str, payload: str, content_type: str, rel_type: str) -> None:
    entries = _parts(path)
    entries[name] = payload.encode("utf-8")
    ct = entries["[Content_Types].xml"].decode("utf-8")
    entries["[Content_Types].xml"] = ct.replace(
        "</Types>", f'<Override PartName="/{name}" ContentType="{content_type}"/></Types>').encode("utf-8")
    rels = entries["xl/_rels/workbook.xml.rels"].decode("utf-8")
    entries["xl/_rels/workbook.xml.rels"] = rels.replace(
        "</Relationships>",
        f'<Relationship Id="rIdCalc" Type="{rel_type}" Target="{name.split("/", 1)[1]}"/></Relationships>',
    ).encode("utf-8")
    buf = io.BytesIO()
    with zipfile.ZipFile(buf, "w", zipfile.ZIP_DEFLATED) as out:
        for n, p in entries.items():
            out.writestr(n, p)
    path.write_bytes(buf.getvalue())


@pytest.fixture()
def book(tmp_path: Path) -> Path:
    wb = openpyxl.Workbook()
    ws = wb.active
    ws.title = "审定表"
    ws["A3"] = "被审计单位："
    ws["A3"].font = Font(bold=True)
    ws["B4"] = "=审定表!A3"
    wb.create_sheet("说明")["A1"] = "不动的页"
    path = tmp_path / "t.xlsx"
    wb.save(path)
    return path


def test_text_patch_keeps_style_and_leaves_other_parts_byte_identical(book: Path):
    before = _parts(book)
    style_before = openpyxl.load_workbook(book).active["A3"].font.b
    out = xp.patch_cells(book, [CellPatch("审定表", "A3", text="被审计单位：某某公司")])
    assert [p.cell for p in out.applied] == ["A3"] and out.skipped == []
    after = _parts(book)
    sheet = _sheet_part(book, "审定表")
    assert out.changed_parts == [sheet]
    assert set(after) == set(before)
    assert all(after[n] == before[n] for n in before if n != sheet)
    ws = openpyxl.load_workbook(book).active
    assert ws["A3"].value == "被审计单位：某某公司" and ws["A3"].font.b is style_before


def test_missing_cell_and_missing_row_are_inserted_in_order(book: Path):
    xp.patch_cells(book, [CellPatch("审定表", "C3", text="右边"), CellPatch("审定表", "A2", text="新行"),
                          CellPatch("审定表", "Z9", text="末行")])
    ws = openpyxl.load_workbook(book).active
    assert (ws["C3"].value, ws["A2"].value, ws["Z9"].value) == ("右边", "新行", "末行")
    xml = _parts(book)[_sheet_part(book, "审定表")].decode("utf-8")
    rows = [int(r) for r in re.findall(r'<row r="(\d+)"', xml)]
    assert rows == sorted(rows)
    assert xml.index('r="A3"') < xml.index('r="C3"')


def test_formula_rewrite_writes_new_cached_value(book: Path):
    xp.patch_cells(book, [CellPatch("审定表", "B4", formula="说明!A1", cached="不动的页")])
    assert openpyxl.load_workbook(book).active["B4"].value == "=说明!A1"
    assert openpyxl.load_workbook(book, data_only=True).active["B4"].value == "不动的页"


def test_cached_only_update_keeps_formula(book: Path):
    xp.patch_cells(book, [CellPatch("审定表", "B4", cached="被审计单位：某某公司")])
    assert openpyxl.load_workbook(book).active["B4"].value == "=审定表!A3"
    assert openpyxl.load_workbook(book, data_only=True).active["B4"].value == "被审计单位：某某公司"


def test_text_with_xml_specials_is_escaped(book: Path):
    xp.patch_cells(book, [CellPatch("审定表", "A3", text="A&B<公司>"), CellPatch("审定表", "B4", cached="X&Y")])
    assert openpyxl.load_workbook(book).active["A3"].value == "A&B<公司>"
    assert openpyxl.load_workbook(book, data_only=True).active["B4"].value == "X&Y"


def test_shared_formula_cells_are_skipped(book: Path):
    part = _sheet_part(book, "审定表")
    _rewrite_part(book, part, lambda x: re.sub(
        r'<c r="B4"[^>]*>.*?</c>',
        '<c r="B4"><f t="shared" ref="B4:B5" si="0">A3</f><v>0</v></c>'
        '<c r="C4"><f t="array" ref="C4">A3</f><v>0</v></c>', x, count=1, flags=re.S))
    before = book.read_bytes()
    out = xp.patch_cells(book, [CellPatch("审定表", "B4", text="x"), CellPatch("审定表", "C4", cached="y")])
    assert out.applied == [] and {p.cell for p, _ in out.skipped} == {"B4", "C4"}
    assert book.read_bytes() == before  # 一格都没改 ⇒ 不落盘


def test_formula_turned_into_text_drops_calc_chain(book: Path):
    _add_part(book, "xl/calcChain.xml",
              '<?xml version="1.0" encoding="UTF-8" standalone="yes"?>'
              '<calcChain xmlns="http://schemas.openxmlformats.org/spreadsheetml/2006/main"><c r="B4" i="1"/></calcChain>',
              "application/vnd.openxmlformats-officedocument.spreadsheetml.calcChain+xml",
              "http://schemas.openxmlformats.org/officeDocument/2006/relationships/calcChain")
    out = xp.patch_cells(book, [CellPatch("审定表", "B4", text="静态")])
    assert out.calc_chain_removed
    parts = _parts(book)
    assert "xl/calcChain.xml" not in parts
    assert b"calcChain" not in parts["[Content_Types].xml"] and b"calcChain" not in parts["xl/_rels/workbook.xml.rels"]
    assert openpyxl.load_workbook(book).active["B4"].value == "静态"


def test_text_patch_without_removed_formula_keeps_calc_chain(book: Path):
    _add_part(book, "xl/calcChain.xml",
              '<?xml version="1.0" encoding="UTF-8" standalone="yes"?>'
              '<calcChain xmlns="http://schemas.openxmlformats.org/spreadsheetml/2006/main"><c r="B4" i="1"/></calcChain>',
              "application/vnd.openxmlformats-officedocument.spreadsheetml.calcChain+xml",
              "http://schemas.openxmlformats.org/officeDocument/2006/relationships/calcChain")
    out = xp.patch_cells(book, [CellPatch("审定表", "A3", text="x")])
    assert not out.calc_chain_removed and "xl/calcChain.xml" in _parts(book)


def test_unknown_sheet_or_bad_coord_raises_and_file_untouched(book: Path):
    before = book.read_bytes()
    with pytest.raises(XlsxPatchError, match="没有页签"):
        xp.patch_cells(book, [CellPatch("不存在", "A1", text="x")])
    with pytest.raises(XlsxPatchError, match="非法坐标"):
        xp.patch_cells(book, [CellPatch("审定表", "A0", text="x")])
    assert book.read_bytes() == before


def _first(pattern: str) -> Path:
    found = sorted(TEMPLATES.rglob(pattern))
    if not found:
        pytest.skip(f"模板 {pattern} 不在（模板目录变了？）")
    return found[0]


def test_append_value_sheets_types_names_and_untouched_parts(book: Path):
    from datetime import date, datetime
    from decimal import Decimal

    before = _parts(book)
    titles = xp.append_value_sheets(book, [
        xp.ValueSheet("明细", [(1, 1, "文字&<>"), (1, 3, 12), (2, 2, 1.5), (3, 1, Decimal("10.20")), (4, 1, True),
                              (5, 1, date(2025, 12, 31)), (6, 1, datetime(2025, 1, 2, 3, 4, 5)),
                              (7, 1, "=不是公式"), (8, 1, "含\x01控制符"), (9, 1, float("nan"))]),
        xp.ValueSheet("审定表", [(1, 1, "重名")]),
        xp.ValueSheet("审定表", [(1, 1, "再重名")]),
    ])
    assert titles == ["明细", "审定表1", "审定表2"]
    after = _parts(book)
    allowed = {"xl/workbook.xml", "xl/_rels/workbook.xml.rels", "[Content_Types].xml", "xl/styles.xml"}
    assert all(after[n] == before[n] for n in before if n not in allowed)
    assert len(after) == len(before) + 3
    wb = openpyxl.load_workbook(book)
    assert wb.sheetnames == ["审定表", "说明", "明细", "审定表1", "审定表2"]
    ws = wb["明细"]
    assert (ws["A1"].value, ws["C1"].value, ws["B2"].value) == ("文字&<>", 12, 1.5)
    assert ws["A3"].value == 10.2 and ws["A4"].value is True
    assert ws["A5"].value == datetime(2025, 12, 31) and ws["A5"].number_format == "mm-dd-yy"
    assert ws["A6"].value == datetime(2025, 1, 2, 3, 4, 5)
    assert ws["A7"].value == "=不是公式" and ws["A7"].data_type == "s", "以 = 开头的文字仍是文字（原先 openpyxl 会当成公式）"
    assert ws["A8"].value == "含控制符" and ws["A9"].value == "nan"
    assert wb["审定表1"]["A1"].value == "重名"


def test_append_value_sheets_without_dates_leaves_styles_untouched(book: Path):
    before = _parts(book)
    xp.append_value_sheets(book, [xp.ValueSheet("明细", [(1, 1, "x")])])
    assert _parts(book)["xl/styles.xml"] == before["xl/styles.xml"]


def _first_title(path: Path) -> str:
    """只读打开会一直占着文件句柄（Windows 上随后的原子替换会被拒），用完即关。"""
    wb = openpyxl.load_workbook(path, read_only=True)
    try:
        return wb.sheetnames[0]
    finally:
        wb.close()


def _dv_count(path: Path) -> int:
    return sum(p.count(b"<x14:dataValidation ") for n, p in _parts(path).items() if n.startswith("xl/worksheets/"))


@pytest.mark.parametrize("pattern", ["C2 销售循环控制测试.xlsx", "H4 工程物资.xlsx"])
def test_real_template_x14_data_validations_survive(tmp_path: Path, pattern: str):
    src = _first(pattern)
    path = tmp_path / src.name
    shutil.copy2(src, path)
    dv = _dv_count(path)
    assert dv > 0, "选的模板应含 x14 数据验证"
    ws_title = _first_title(path)
    before = _parts(path)
    out = xp.patch_cells(path, [CellPatch(ws_title, "A2", text="被审计单位：某某公司")])
    after = _parts(path)
    assert _dv_count(path) == dv
    assert all(after[n] == before[n] for n in before if n not in out.changed_parts)


def test_real_xlsm_macros_survive(tmp_path: Path):
    src = _first("B23-1 *.xlsm")
    path = tmp_path / src.name
    shutil.copy2(src, path)
    vba = hashlib.sha256(_parts(path)["xl/vbaProject.bin"]).hexdigest()
    title = _first_title(path)
    xp.patch_cells(path, [CellPatch(title, "A2", text="被审计单位：某某公司")])
    assert hashlib.sha256(_parts(path)["xl/vbaProject.bin"]).hexdigest() == vba
    assert openpyxl.load_workbook(path, keep_vba=True)[title]["A2"].value == "被审计单位：某某公司"

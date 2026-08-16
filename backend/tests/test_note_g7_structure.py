"""附注 G7 长期股权投资章节结构守卫（四作用域：上市 五、18 / 七、1，国企 八、18 / 七、*13节）。

锁定四个幂等脚本的对齐结果，用 openpyxl 直读源 xlsx 与模板 `headers`/`columns`/`rows`
做**三向比对**，并覆盖：

1. 两级表头 `本期增减变动` 分组的列索引与 `_column_groups` 一致；
2. 上市联营 FS 表行集为 17 行且**不含**「其中：现金和现金等价物」
   （合营表 18 行才有该行 —— 源模板 R135 只在合营表出现）；
3. 国企两张历史 `续：` 表已正名为不同表名（不再同名互相覆盖）；
4. 国企两张段落文本泄漏表名已正名为可读表名；
5. 反向自检：压扁的 5 列 headers（脚本改造前状态）必须被 `validate_section` 判红。

spec: .kiro/specs/g7-four-table-extraction-and-disclosure-alignment/ (Task 4.5)
"""
from __future__ import annotations

import importlib.util
import json
import sys
from pathlib import Path

import openpyxl
import pytest

_ROOT = Path(__file__).resolve().parent.parent
_FIX_DIR = _ROOT / "scripts" / "fix"
_SRC_XLSX = _ROOT / "wp_templates" / "G" / "G7 长期股权投资.xlsx"
_LISTED = _ROOT / "data" / "note_template_listed.json"
_SOE = _ROOT / "data" / "note_template_soe.json"

sys.path.insert(0, str(_FIX_DIR))


def _load_fix(name: str):
    spec = importlib.util.spec_from_file_location(f"_fix_{name}", _FIX_DIR / f"{name}.py")
    mod = importlib.util.module_from_spec(spec)
    spec.loader.exec_module(mod)
    return mod


FIX_MAIN = _load_fix("fix_note_g7_long_term_equity_structure")
FIX_LISTED_OTHER = _load_fix("fix_note_g7_listed_other_entities_structure")
FIX_SOE = _load_fix("fix_note_g7_soe_structure")
FIX_SCOPE = _load_fix("fix_note_g7_soe_scope_change_structure")


def _section(path: Path, section_number: str) -> dict:
    doc = json.loads(path.read_text(encoding="utf-8"))
    return next(s for s in doc["sections"] if str(s.get("section_number")) == section_number)


# ─────────────────────────── --check 全绿（四脚本）───────────────────────────


class TestAllScriptsCheckPass:
    def test_listed_main(self):
        _c, w, errs = FIX_MAIN._runner("listed", dry_run=True, check=True)
        assert not errs, errs
        assert not w, w

    def test_listed_other_entities(self):
        _c, w, errs = FIX_LISTED_OTHER._runner("listed", dry_run=True, check=True)
        assert not errs, errs
        assert not w, w

    def test_soe_main(self):
        _c, w, errs = FIX_SOE._runner("soe", dry_run=True, check=True)
        assert not errs, errs
        assert not w, w

    @pytest.mark.parametrize("key", list(FIX_SCOPE._plan().keys()))
    def test_soe_scope_change(self, key):
        _c, w, errs = FIX_SCOPE._runner(key, dry_run=True, check=True)
        assert not errs, (key, errs)
        assert not w, (key, w)


# ─────────────────────────── 上市 五、18 主表 ───────────────────────────


class TestListedMainTable:
    def test_two_level_header_group(self):
        sec = _section(_LISTED, "五、18")
        tbl = (sec.get("tables") or [])[0]
        cols = tbl.get("columns") or []
        # 🔴 标签列 key 于 g7-column-alignment spec Task 5 由 '项目' 统一为平台惯例 'label'
        # （平台 266 个标签列定义里 241 个用 'label'，跨 70 文件；'项目' 全平台仅 7 处且全在
        # G 循环 = 少数派偏离 + 中文字面量当 key）。此处原断言锁死的是 Task 5 之前的状态，
        # 已按落地后事实诚实改写；下方 test_label_key_is_platform_convention 反向锁死防回退。
        assert [c["key"] for c in cols] == [
            "label", "openingBook", "openingImpairment", "addition", "reduction",
            "equityProfit", "oci", "otherEquity", "dividend", "impairment", "other",
            "closingBook", "closingImpairment",
        ]
        groups = tbl.get("_column_groups")
        assert groups == [{"group": "本期增减变动", "start": 3, "span": 8}], groups

    def test_label_key_is_platform_convention(self):
        """反向锁死：标签列 key 必须是 'label'，且不得回退成中文字面量当 key。

        改回 '项目'（或任何中文 key）即打红 —— 这是 Task 5 的防回退断言。
        `is_label` / `label` 显示文字不在本断言范围内（前者是投影器选标签列的依据、
        后者是源 xlsx 原文，两者都不因 key 统一而变）。
        """
        sec = _section(_LISTED, "五、18")
        tbl = (sec.get("tables") or [])[0]
        cols = tbl.get("columns") or []
        label_cols = [c for c in cols if c.get("is_label")]
        assert len(label_cols) == 1, label_cols
        assert label_cols[0]["key"] == "label", label_cols[0]
        # 显示文字仍取源 xlsx 原文，未被 key 统一波及
        assert label_cols[0]["label"] == "被投资单位", label_cols[0]
        # 全表任何列的 key 都不得是中文字面量
        cjk_keys = [c["key"] for c in cols if any("\u4e00" <= ch <= "\u9fff" for ch in c["key"])]
        assert cjk_keys == [], f"列 key 不得用中文字面量（违反禁硬编码）: {cjk_keys}"

    def test_headers_match_source_xlsx(self):
        """源模板第 6 列「权益法下确认的投资损益」跨 R9/R10/R11 三行拼字，
        第 7 列「其他综合收益调整」跨 R9/R10 两行拼字；其余单行列取 R9（或跨列合并的 R8）。"""
        wb = openpyxl.load_workbook(_SRC_XLSX, data_only=True)
        ws = wb["附注披露信息（上市公司）"]

        def col_text(col: int, rows: tuple[int, ...]) -> str:
            return "".join(str(ws.cell(r, col).value or "") for r in rows)

        expected_flat = [
            str(ws.cell(8, 1).value),  # 被投资单位
            str(ws.cell(8, 2).value),  # 期初余额（账面价值）
            str(ws.cell(9, 3).value),  # 减值准备期初余额
            col_text(4, (9,)), col_text(5, (9,)),
            col_text(6, (9, 10, 11)),  # 权益法下确认的投资损益
            col_text(7, (9, 10)),      # 其他综合收益调整
            col_text(8, (9,)), col_text(9, (9,)), col_text(10, (9,)), col_text(11, (9,)),
            str(ws.cell(8, 12).value),  # 期末余额（账面价值）
            str(ws.cell(8, 13).value),  # 减值准备期末余额
        ]
        tbl = (_section(_LISTED, "五、18").get("tables") or [])[0]
        assert tbl["headers"] == expected_flat

    def test_row_set_no_header_label(self):
        tbl = (_section(_LISTED, "五、18").get("tables") or [])[0]
        rows = tbl.get("rows") or []
        assert [r["label"] for r in rows] == [
            "①合营企业", "…", "小计", "②联营企业", "…", "小计", "合计",
        ]
        assert not any(r.get("row_type") == "header_label" for r in rows)

    def test_no_html_in_headers(self):
        tbl = (_section(_LISTED, "五、18").get("tables") or [])[0]
        assert not any("<" in h for h in tbl["headers"])


# ─────────────────────────── 上市 七、1（14 表）───────────────────────────


class TestListedOtherEntities:
    EXPECTED_TABLES = [
        "企业集团的构成", "重要的非全资子公司", "重要非全资子公司主要财务信息—期末数",
        "续（1）", "续（2）", "未丧失控制权的所有者权益份额变动影响",
        "重要的合营企业或联营企业", "重要合营企业主要财务信息",
        "续：重要合营企业本期及上期经营成果", "重要联营企业主要财务信息",
        "续：重要联营企业本期及上期经营成果",
        "其他不重要合营企业和联营企业的汇总财务信息",
        "对合营企业或联营企业发生超额亏损的分担额", "重要的共同经营",
    ]

    def test_fourteen_tables_no_orphans(self):
        sec = _section(_LISTED, "七、1")
        names = [t.get("name") for t in (sec.get("tables") or [])]
        assert names == self.EXPECTED_TABLES

    def test_associate_fs_seventeen_rows_no_cash_row(self):
        """🔴 Property 12：联营 FS 表 17 行且不含现金及现金等价物行；合营表 18 行含该行。"""
        sec = _section(_LISTED, "七、1")
        by_name = {t["name"]: t for t in sec["tables"]}

        jv_rows = [r["label"] for r in by_name["重要合营企业主要财务信息"]["rows"]]
        assert len(jv_rows) == 18
        assert "其中：现金和现金等价物" in jv_rows

        assoc_rows = [r["label"] for r in by_name["重要联营企业主要财务信息"]["rows"]]
        assert len(assoc_rows) == 17
        assert "其中：现金和现金等价物" not in assoc_rows

    def test_all_tables_have_columns_and_guidance(self):
        sec = _section(_LISTED, "七、1")
        for t in sec["tables"]:
            assert t.get("columns"), t["name"]
            assert str(t.get("guidance") or "").strip(), t["name"]

    def test_text_sections_no_bare_examples(self):
        """真实需要填写的说明段（剔除示例数字与【提示】括注）。"""
        sec = _section(_LISTED, "七、1")
        texts = sec.get("text_sections") or []
        assert len(texts) == 20  # 10 组 (#### 标题 + 正文)
        joined = "\n".join(texts)
        assert "【" not in joined, "【提示】括注不应作为正式披露正文"
        assert "200Y" not in joined and "2012年10月" not in joined, "示例数字不应留在正文"


# ─────────────────────────── 国企 八、18（10 表）───────────────────────────


class TestSoeMainSection:
    EXPECTED_TABLES = [
        "长期股权投资分类", "长期股权投资明细",
        "重要合营企业的主要财务信息（划分为持有待售的除外）",
        "续：重要合营企业本期及上期经营成果",
        "重要联营企业的主要财务信息",
        "续：重要联营企业本期及上期经营成果",
        "不重要合营企业和联营企业的汇总信息",
        "②对合营企业或联营企业发生超额亏损的分担额",
        "结构化主体权益的账面价值和最大损失敞口",
        "结构化主体获得收益及转移资产情况",
    ]

    def test_ten_tables_no_duplicate_names(self):
        """🔴 两张历史同名『续：』已正名为不同表名（不再互相覆盖丢表）。"""
        sec = _section(_SOE, "八、18")
        names = [t.get("name") for t in (sec.get("tables") or [])]
        assert names == self.EXPECTED_TABLES
        assert len(names) == len(set(names)), "表名仍有重复"
        assert "续：" not in names

    def test_orphan_paragraph_names_renamed(self):
        """两张段落文本泄漏表名已正名为可读表名。"""
        sec = _section(_SOE, "八、18")
        names = {t.get("name") for t in sec["tables"]}
        assert "结构化主体权益的账面价值和最大损失敞口" in names
        assert "结构化主体获得收益及转移资产情况" in names
        for orphan in (
            "C.在财务报表中确认的与企业在未纳入合并财务报表范围的结构化主体中"
            "权益相关的资产和负债的账面价值与其最大损失敞口的比较。",
            "本公司发起多个结构化主体，但在结构化中均不持有权益。",
        ):
            assert not any(orphan in (n or "") for n in names), names

    def test_classification_table_has_six_rows(self):
        """长期股权投资分类补齐源模板 R203「对子公司投资」行（原 5 行漏一行）。"""
        sec = _section(_SOE, "八、18")
        tbl = next(t for t in sec["tables"] if t["name"] == "长期股权投资分类")
        labels = [r["label"] for r in tbl["rows"]]
        assert labels == [
            "对子公司投资", "对合营企业投资", "对联营企业投资",
            "小  计", "减：长期股权投资减值准备", "合  计",
        ]

    def test_detail_table_two_level_header(self):
        sec = _section(_SOE, "八、18")
        tbl = next(t for t in sec["tables"] if t["name"] == "长期股权投资明细")
        groups = tbl.get("_column_groups")
        assert groups == [{"group": "本期增减变动", "start": 3, "span": 8}], groups

    def test_associate_fs_uses_corrected_label(self):
        """R268 字面「对合营企业权益投资的账面价值」处于联营表内为源模板笔误，
        沿用附注模板已修正口径「对联营企业权益投资的账面价值」（design.md R5.6）。
        """
        sec = _section(_SOE, "八、18")
        tbl = next(t for t in sec["tables"] if t["name"] == "重要联营企业的主要财务信息")
        labels = [r["label"] for r in tbl["rows"]]
        assert "对联营企业权益投资的账面价值 " in labels
        assert "对合营企业权益投资的账面价值 " not in labels

    def test_all_tables_have_columns_and_guidance(self):
        sec = _section(_SOE, "八、18")
        for t in sec["tables"]:
            assert t.get("columns"), t["name"]
            assert str(t.get("guidance") or "").strip(), t["name"]
            for h in t.get("headers") or []:
                assert "<" not in h, (t["name"], h)


# ─────────────────────────── 国企 七、* 13 节 ───────────────────────────


class TestSoeScopeChangeSections:
    def test_thirteen_sections_covered(self):
        plan = FIX_SCOPE._plan()
        assert len(plan) == 13

    def test_no_bare_header_row_names(self):
        """表头首格泄漏名（序号/公司名称）已正名。"""
        plan = FIX_SCOPE._plan()
        for section_number, (_label, _plan_rules, expected, _txt) in plan.items():
            sec = _section(_SOE, section_number)
            names = [t.get("name") for t in (sec.get("tables") or [])]
            for n in names:
                assert n not in ("序号", "公司名称"), (section_number, names)

    def test_sale_date_naming_not_disposal_date(self):
        """🔴 源模板逐字为「出售日」，模板 JSON 曾漂移写成「处置日」，本脚本已正名。"""
        sec = _section(_SOE, "七、本期不再纳入合并")
        names = [t.get("name") for t in sec["tables"]]
        assert "本期出售的子公司出售日的经营成果" in names
        assert "本期出售的子公司处置日的经营成果" not in names

    def test_text_only_sections_have_required_text(self):
        """两个纯文本节（无表格）补了源模板要求的说明段。"""
        for section_number in ("七、子公司使用企业集", "七、纳入合并财务报表"):
            sec = _section(_SOE, section_number)
            assert not sec.get("tables")
            assert sec.get("text_sections"), section_number

    def test_all_declared_tables_have_columns(self):
        plan = FIX_SCOPE._plan()
        for section_number, (_label, _rules, expected, _txt) in plan.items():
            if not expected:
                continue
            sec = _section(_SOE, section_number)
            by_name = {t["name"]: t for t in sec["tables"]}
            for name in expected:
                assert by_name[name].get("columns"), (section_number, name)
                assert str(by_name[name].get("guidance") or "").strip(), (section_number, name)


# ─────────────────────────── 反向自检 ───────────────────────────


class TestReverseSelfCheck:
    def test_validate_section_catches_flattened_headers(self):
        """反向自检：改造前状态（压扁 5 列、无 columns）必须被判红，证明守卫真的在检查。"""
        from _note_structure_kit import validate_section

        broken = {
            "tables": [
                {
                    "name": "长期股权投资",
                    "headers": [
                        "被投资单位", "期初余额（账面价值）", "本期增减变动",
                        "期末余额（账面价值）", "减值准备期末余额",
                    ],
                    "rows": [{"label": "被投资单位", "row_type": "header_label"}],
                    "guidance": None,
                }
            ]
        }
        errs = validate_section(broken, ["长期股权投资"])
        assert any("缺 columns" in e for e in errs), errs
        assert any("header_label" in e for e in errs), errs

    def test_validate_section_catches_duplicate_table_names(self):
        from _note_structure_kit import validate_section

        broken = {
            "tables": [
                {"name": "续：", "headers": ["项目"], "columns": [{"key": "item", "label": "项目", "is_label": True, "flat": True}], "rows": [], "guidance": "x"},
                {"name": "续：", "headers": ["项目"], "columns": [{"key": "item", "label": "项目", "is_label": True, "flat": True}], "rows": [], "guidance": "x"},
            ]
        }
        errs = validate_section(broken, ["续："])
        assert any("重复" in e for e in errs), errs


# ─────────────── `--check` 判据强度（平台级防回退，2026-08-12）───────────────


class TestCheckModeStrength:
    """`_note_structure_kit.run_section` 的 `--check` 必须与**写入路径**同源。

    ## 为什么这条守卫必须存在

    改造前 `--check` 分支只跑 `validate_section`（表名清单 + `text_sections`），
    **完全不跑 `apply_plan`** ⇒ `headers` / `columns` / `rows` / `guidance` /
    `_column_groups` 的偏差它一律看不见。实测形态：

        fix_note_g7_soe_structure.py --dry-run   → 共 2 处变更
        fix_note_g7_soe_structure.py --check     → 0 项欠账      ← 同一时刻

    而各 spec（含本仓 G7 spec 的「幂等双证」任务）普遍拿「`--check` 0 欠账」
    当对齐/幂等的验收判据 ⇒ **验收恒过、列结构欠账长期驻留**。这是平台级假绿，
    且 24 个幂等脚本共用这一条通路，故守卫钉在共享 kit 上而不是逐脚本重复。

    本类的判据是**行为级**（真跑一次 check、看偏差是否被计为欠账），
    不是「源码里是否出现 apply_plan 字样」—— 后者被 `if False:` 之类一改就假绿。
    """

    @staticmethod
    def _kit():
        import sys as _sys

        _sys.path.insert(0, str(_ROOT / "scripts" / "fix"))
        import _note_structure_kit as kit

        return kit

    @classmethod
    def _one_table_doc(cls, label_header: str, guidance: str = "g") -> dict:
        """最小可用章节：1 张表、1 个标签列 + 1 个数据列。

        🔴 `columns` / `rows` 必须由 **kit 自己的构造器**生成（不手写字面量）——
        手写版少一个字段（如 `format: None`）就会让「已对齐」的替身也报变更，
        那是**替身缺陷**，会把本类的反面用例（对齐时应 0 欠账）打成假红。
        """
        kit = cls._kit()
        return {
            "sections": [
                {
                    "section_number": "测试、1",
                    "tables": [
                        {
                            "name": "T",
                            "headers": [label_header, "金额"],
                            "columns": kit.flat_columns(
                                [("label", label_header, None), ("amt", "金额", kit.AMOUNT)]
                            ),
                            "rows": [kit.data_row("行1")],
                            "guidance": guidance,
                        }
                    ],
                }
            ]
        }

    def _run_check(self, kit, tmp_path, disk_label: str, want_label: str):
        """把 `disk_label` 写盘、用 `want_label` 作目标态跑 `--check`，返回 errs。"""
        import json as _json

        path = tmp_path / "note_template_probe.json"
        path.write_text(
            _json.dumps(self._one_table_doc(disk_label), ensure_ascii=False),
            encoding="utf-8",
        )
        plan = [
            kit.rule(
                "T",
                kit.flat_columns([("label", want_label, None), ("amt", "金额", kit.AMOUNT)]),
                [kit.data_row("行1")],
                "g",
            )
        ]
        changes, warnings, errs = kit.run_section(
            path,
            "测试、1",
            plan,
            ["T"],
            aligned_by="probe",
            dry_run=False,
            check=True,
        )
        # check 模式绝不写盘
        assert _json.loads(path.read_text(encoding="utf-8"))["sections"][0]["tables"][0][
            "headers"
        ][0] == disk_label, "check 模式写盘了 —— 只读性被破坏"
        assert changes == [] and warnings == [], (changes, warnings)
        return errs

    def test_check_catches_column_label_drift(self, tmp_path):
        """列 label 偏差必须被 `--check` 计为欠账（改造前此处恒 0 欠账）。"""
        kit = self._kit()
        errs = self._run_check(kit, tmp_path, disk_label="项目", want_label="项  目")
        assert errs, "列 label 偏差未被 --check 抓到 ⇒ 假绿回来了"
        assert any("结构未对齐" in e for e in errs), errs

    def test_check_is_green_when_aligned(self, tmp_path):
        """反面：磁盘与目标态一致时必须 0 欠账（防判据写成恒红）。"""
        kit = self._kit()
        errs = self._run_check(kit, tmp_path, disk_label="项  目", want_label="项  目")
        assert errs == [], errs

    def test_check_and_dry_run_agree_on_same_input(self, tmp_path):
        """同一输入下 `--check` 有欠账 ⟺ `--dry-run` 有变更（两口径不得再分叉）。"""
        import json as _json

        kit = self._kit()
        for disk_label, aligned in (("项目", False), ("项  目", True)):
            path = tmp_path / f"probe_{aligned}.json"
            path.write_text(
                _json.dumps(self._one_table_doc(disk_label), ensure_ascii=False),
                encoding="utf-8",
            )
            plan = [
                kit.rule(
                    "T",
                    kit.flat_columns([("label", "项  目", None), ("amt", "金额", kit.AMOUNT)]),
                    [kit.data_row("行1")],
                    "g",
                )
            ]
            _c, _w, errs = kit.run_section(
                path, "测试、1", plan, ["T"], aligned_by="probe", dry_run=False, check=True
            )
            changes, _w2, _e2 = kit.run_section(
                path, "测试、1", plan, ["T"], aligned_by="probe", dry_run=True, check=False
            )
            assert bool(errs) == bool(changes), (
                f"disk_label={disk_label!r}: --check 欠账 {len(errs)} 处 "
                f"vs --dry-run 变更 {len(changes)} 处 —— 两口径分叉即假绿复发"
            )
            assert bool(changes) is not aligned

    def test_check_source_replays_write_path(self):
        """源码级辅助判据：check 分支必须真调 `apply_plan`（在深拷贝上）。

        单独看源码不足以证明判据有效（`if False:` 也能留住字样），故本条只作
        **补充**；行为级判据是上面三条。这里额外钉住「深拷贝」——
        少了它 `--check` 会污染内存态、并可能被后续逻辑写盘。
        """
        src = (_ROOT / "scripts" / "fix" / "_note_structure_kit.py").read_text(encoding="utf-8")
        head = src.split("if check:", 1)
        assert len(head) == 2, "check 分支锚点未命中"
        branch = head[1].split("return [], [], errs", 1)[0]
        assert "copy.deepcopy(section)" in branch, "check 分支缺深拷贝"
        assert "apply_plan(probe" in branch, "check 分支未在副本上重放 apply_plan"

"""形态 B 的 `sheet` 三分支守卫 —— `l2` / `m1` 六个既有 handler（additive 兼容）。

spec: `x3-adjustment-entry-import-export` · Wave 3 任务 5.2
_Requirements: 4.1, 4.2, 4.3, 4.4, 1.6, 11.2_

## 判据为什么不连库

任务 5.2 的施加物是「六个既有 handler 追加 `sheet: str | None = Query(None)`，并按值三分支」。
需要证明的是**分支选择**与**既有分支不受影响**，不是数据库往返（那是任务 15.1 的作业面）：

- `sheet` 不传 ⇒ FastAPI 落 `Query(None)` 的默认值 ⇒ `None` ⇒ 走既有路径。
  「不传等价于 None」这一步由 `test_sheet_param_is_optional_with_none_default` 从**运行期
  已解析的 `dependant.query_params`** 取，不靠读签名文本（父 spec 踩过 `getattr(field,
  "required", False)` 在 `fastapi._compat.v2.ModelField` 上不存在、fail-open 判成非必填）。
- 「既有分支不受影响」由 `test_legacy_body_never_reads_sheet` 用 **AST** 证明：六个 handler
  里 `sheet` 只出现在新增的那一句守卫上，既有语句一处都没读它 ⇒ 守卫为假时，剩下的代码
  在结构上不可能因 `sheet` 而变。这比「跑一遍看结果一样」强：后者只覆盖被跑到的那条路径。
- 唯一不依赖 DB 的 handler（`l2` 的 export-template）另做一次**真实调用**，比对产物结构与
  既有构造函数逐字一致 ⇒ 证明 `None` 分支真的落在既有代码上、没有悄悄走委派。

## 最贵的那条判据：前端在传什么

`useL2ImportExport` / `useM1ImportExport` **一直在传** `sheet`（4 个组件共 12 处调用），
而施加前六个 handler 没有该参数 ⇒ FastAPI 静默丢弃（E5/E6 存量缺陷）。施加后
「值 ∉ `IE_SHEETS` ⇒ 400」这条一旦把这些值挡在外面，线上这 4 个页面的导入导出会**全废**
（接口 400，而后端单测、Volar、vitest 三层都不会红）。故本文件从**前端源码**抽出实际传的
sheet 取值，双向锁死它们仍走既有分支 —— 这一条是唯一能机械抓到该回归的判据。
"""

from __future__ import annotations

import ast
import inspect
import io
import re
from pathlib import Path
from types import ModuleType
from typing import Any

import pytest
from fastapi import HTTPException
from openpyxl import load_workbook

from app.routers import l2_interest_payable as l2_host
from app.routers import m1_dividends_payable as m1_host

_REPO = Path(__file__).resolve().parents[2]
_WP_ROOT = _REPO / "audit-platform/frontend/src/components/workpaper"

_SHARED_IMPL = "app.routers.wp_render_strategies._x3_adjustment_import_export"

#: 三态后缀（与 design §C1 / 前端共享 composable 同一组字面量）
_ACTIONS: tuple[str, ...] = ("export-template", "export-data", "import-data")

#: 短前缀 → (宿主模块, 长前缀, 六个 handler 的函数名)
_HOSTS: dict[str, tuple[ModuleType, str, dict[str, str]]] = {
    "l2": (
        l2_host,
        "l2-interest-payable",
        {
            "export-template": "l2_export_template",
            "export-data": "l2_export_data",
            "import-data": "l2_import_data",
        },
    ),
    "m1": (
        m1_host,
        "m1-dividends-payable",
        {
            "export-template": "m1_export_template",
            "export-data": "m1_export_data",
            "import-data": "m1_import_data",
        },
    ),
}

#: 前端实际调用这些 handler 的组件（`useL2ImportExport` / `useM1ImportExport` 的消费方）。
#: 🔴 只许增不许减：删一个就少锁一页，那一页被 400 挡住时没人会红。
_FRONTEND_CALLERS: dict[str, tuple[str, ...]] = {
    "l2": ("l2/core/L2TabDetail.vue",),
    "m1": (
        "m1/core/M1TabDetail.vue",
        "m1/calc/M1TabFxRate.vue",
        "m1/calc/M1TabDividendCalc.vue",
    ),
}

#: `exportTemplate('M1-2')` / `exportData('M1-4')` / `importData(file, 'M1-5')` 三种调用形态
_FE_CALL_RE = re.compile(
    r"(?:exportTemplate|exportData|importData)\(\s*(?:file\s*,\s*)?'([^']+)'"
)


def _legacy_route(short: str, action: str) -> Any:
    """取形态 B（长前缀）那条路由对象。"""
    module, long_prefix, _ = _HOSTS[short]
    want = f"/api/{long_prefix}/{{wp_id}}/{action}"
    hits = [r for r in module.router.routes if getattr(r, "path", "") == want]
    assert len(hits) == 1, f"{want} 命中 {len(hits)} 条（期望 1）—— 形态 B 端点被改动"
    return hits[0]


def _shape_a_route(short: str, action: str) -> Any:
    module, _, _ = _HOSTS[short]
    want = f"/api/workpapers/{{wp_id}}/{short}/{action}"
    hits = [r for r in module.router.routes if getattr(r, "path", "") == want]
    assert len(hits) == 1, f"{want} 命中 {len(hits)} 条（期望 1）—— 形态 A 未挂载"
    return hits[0]


def _frontend_sheet_args(short: str) -> dict[str, list[str]]:
    """从前端调用点抽出实际传给这三个端点的 sheet 取值。"""
    out: dict[str, list[str]] = {}
    for rel in _FRONTEND_CALLERS[short]:
        path = _WP_ROOT / rel
        assert path.is_file(), f"前端调用点缺失：{path}（守卫必须打红而非跳过）"
        out[rel] = _FE_CALL_RE.findall(path.read_text(encoding="utf-8"))
    return out


@pytest.mark.parametrize("short", sorted(_HOSTS))
@pytest.mark.parametrize("action", _ACTIONS)
def test_sheet_param_is_optional_with_none_default(short: str, action: str) -> None:
    """六个既有 handler 各声明 `sheet`，且**非必填、默认 None**（additive，R4.3 / R1.6）。

    判据取 FastAPI 已解析的 `dependant.query_params[].field_info`：
    - 必填了 ⇒ 既有调用方不传就 422 = 破坏性变更；
    - 压根没声明 ⇒ 前端传的值被静默丢弃（施加前的存量缺陷形态）；
    - 判不出来（`field_info` 无 `is_required`）⇒ 本条**失败**，不返回「看起来没问题」。
    """
    route = _legacy_route(short, action)
    fields = {
        f.name: f for f in (getattr(route.dependant, "query_params", []) or [])
    }
    assert "sheet" in fields, (
        f"{short}/{action} 未声明 sheet query 参数 —— FastAPI 会静默丢弃前端传的值，"
        "X-3 页导出到的是业务表内容（200 + 错文件）"
    )
    info = fields["sheet"].field_info
    assert hasattr(info, "is_required"), (
        f"{short}/{action}: field_info={type(info).__name__} 无 is_required —— 判据失效"
    )
    assert info.is_required() is False, (
        f"{short}/{action}: sheet 是必填 —— 既有调用方不传会 422（违反 R1.6 additive 要求）"
    )
    assert info.default is None, (
        f"{short}/{action}: sheet 默认值 {info.default!r} != None —— "
        "不传时会落到一个业务默认值上，那是「串表导出」的形态"
    )


@pytest.mark.parametrize("short", sorted(_HOSTS))
def test_none_takes_legacy_branch(short: str) -> None:
    """`sheet=None`（= 不传）⇒ 判定为「不委派」⇒ 走既有路径（R1.6）。"""
    module, _, _ = _HOSTS[short]
    assert module._requires_x3(None) is False, (
        f"{short}: sheet=None 被判成需要委派 —— 既有调用方（不传 sheet）会被改道，"
        "这是 additive 的反面"
    )


@pytest.mark.parametrize("short", sorted(_HOSTS))
def test_legacy_sheet_codes_take_legacy_branch(short: str) -> None:
    """既有业务表码 ⇒ 不委派、不 400，产物与施加前逐字节相同（R1.6）。"""
    module, _, _ = _HOSTS[short]
    legacy = sorted(module.IE_SHEETS - module._X3_CODES)
    assert legacy, f"{short}: IE_SHEETS 去掉 X-3 后为空 —— 既有可服务面被抹掉了"
    for code in legacy:
        assert module._requires_x3(code) is False, (
            f"{short}: sheet={code} 被判成需要委派 —— 既有业务表会被导成 X-3 内容"
        )


@pytest.mark.parametrize("short", sorted(_HOSTS))
def test_frontend_passed_sheets_are_whitelisted(short: str) -> None:
    """前端在传的 sheet 取值必须 ∈ `IE_SHEETS` 且走既有分支（本文件最贵的一条）。

    施加前这些值被 FastAPI 静默丢弃；施加后它们要么走既有分支（正确），要么撞上
    「∉ `IE_SHEETS` ⇒ 400」⇒ 线上该页导入导出全废。**只有这条能机械抓到**：
    后端其余判据都只看后端自己声明的白名单，看不见前端在传什么。
    """
    module, _, _ = _HOSTS[short]
    per_file = _frontend_sheet_args(short)
    for rel, args in per_file.items():
        assert len(args) >= len(_ACTIONS), (
            f"{rel} 只抽出 {len(args)} 个 sheet 实参（期望 ≥ {len(_ACTIONS)} = 三态各一）"
            " —— 抽取正则失效或调用点被改，本条已空转"
        )
    passed = sorted({code for args in per_file.values() for code in args})
    unknown = [c for c in passed if c not in module.IE_SHEETS]
    assert not unknown, (
        f"{short}: 前端在传但不在 IE_SHEETS 里的 sheet：{unknown}"
        f"（IE_SHEETS={sorted(module.IE_SHEETS)}）\n"
        "→ 这些请求会撞上 400 分支，线上该页导入导出全废。正解是把它们纳入白名单真源"
        "（`l2` 从 `_SHEET_NAME` 派生 / `m1` 从 service 的 `EXPORT_SHEETS` 派生），"
        "不是在前端把参数删掉。"
    )
    delegated = [c for c in passed if module._requires_x3(c)]
    assert not delegated, (
        f"{short}: 前端既有页面传的 {delegated} 被判成需要委派共享实现 —— "
        "业务表页会导出 X-3 的内容"
    )


@pytest.mark.parametrize("short", sorted(_HOSTS))
def test_x3_code_takes_delegate_branch(short: str) -> None:
    """X-3 码 ⇒ 委派，且委派目标就是形态 A 那三个端点的**同一个** callable（R1.5）。"""
    module, _, _ = _HOSTS[short]
    codes = sorted(module._X3_CODES)
    assert len(codes) == 1, f"{short}: 反查到 {codes}（期望恰 1 张 X-3）"
    assert module._requires_x3(codes[0]) is True, (
        f"{short}: sheet={codes[0]} 未被判成委派 —— X-3 会落到既有业务表代码路径（串表导出）"
    )
    for action in _ACTIONS:
        route = _shape_a_route(short, action)
        assert module._X3_SHAPE_A[action] is route.endpoint, (
            f"{short}/{action}: 委派目标不是形态 A 的端点对象 —— 抄了第二套编排（R1.5）"
        )
        assert getattr(route.endpoint, "__module__", "?") == _SHARED_IMPL, (
            f"{short}/{action}: 形态 A 端点宿主 "
            f"{getattr(route.endpoint, '__module__', '?')} != 共享实现"
        )


@pytest.mark.parametrize("short", sorted(_HOSTS))
def test_unregistered_sheet_raises_readable_400(short: str) -> None:
    """未登记 `sheet` ⇒ 400 可读错误，**不得**回退成「导出全部 sheet」（R4.2）。

    回退是 E6 存量缺陷的形态：200 + 错内容，比 400 难查一个量级。故这里同时钉住
    「不是返回 False」——返回 False 就等于悄悄走既有整表路径。
    """
    module, _, _ = _HOSTS[short]
    others = sorted(
        c
        for other, (mod, _, _) in _HOSTS.items()
        if other != short
        for c in mod.IE_SHEETS
    )
    for bad in ["ZZ-9", "", "l2-3", *others]:
        with pytest.raises(HTTPException) as exc:
            module._requires_x3(bad)
        assert exc.value.status_code == 400, (
            f"{short}: sheet={bad!r} 得到 {exc.value.status_code}（期望 400）"
        )
        detail = str(exc.value.detail)
        assert bad in detail or not bad, f"{short}: 400 文案未回显收到的取值：{detail}"
        for code in sorted(module.IE_SHEETS):
            assert code in detail, (
                f"{short}: 400 文案未列出支持的 sheet {code}：{detail}"
            )


@pytest.mark.parametrize("short", sorted(_HOSTS))
@pytest.mark.parametrize("action", _ACTIONS)
def test_legacy_body_never_reads_sheet(short: str, action: str) -> None:
    """AST：`sheet` 只出现在新增的那句守卫里，既有语句一处都没读它（R1.6 的结构证明）。

    守卫为假时剩下的代码在结构上不可能因 `sheet` 而改变行为 —— 这比「跑一遍结果一样」强，
    后者只覆盖被跑到的那一条路径。
    """
    module, _, names = _HOSTS[short]
    func = getattr(module, names[action])
    tree = ast.parse(inspect.cleandoc(inspect.getsource(func)))
    body = tree.body[0].body  # type: ignore[attr-defined]
    stmts = [s for s in body if not isinstance(s, ast.Expr)]  # 跳过 docstring
    assert stmts, f"{short}/{action}: handler 主体为空 —— 取源码失败"

    guard = stmts[0]
    assert isinstance(guard, ast.If), (
        f"{short}/{action}: 首条语句不是 `if _requires_x3(sheet):` 守卫，实为 "
        f"{type(guard).__name__} —— 分派点必须在既有语句**之前**，否则既有副作用"
        "（如 `await file.read()`）会先发生"
    )
    guard_src = ast.dump(guard.test)
    assert "_requires_x3" in guard_src and "'sheet'" in guard_src, (
        f"{short}/{action}: 守卫条件不是 `_requires_x3(sheet)`：{ast.unparse(guard.test)}"
    )
    assert len(guard.body) == 1 and isinstance(guard.body[0], ast.Return), (
        f"{short}/{action}: 守卫体不是单条 return（委派应直接返回，不与既有路径混跑）"
    )
    assert not guard.orelse, f"{short}/{action}: 守卫带 else 分支 —— 既有路径应在守卫之外"

    leaked = [
        node
        for stmt in stmts[1:]
        for node in ast.walk(stmt)
        if isinstance(node, ast.Name) and node.id == "sheet"
    ]
    assert not leaked, (
        f"{short}/{action}: 既有语句里出现了 {len(leaked)} 处 `sheet` 引用 —— "
        "既有行为已被新参数影响，不再是 additive（R1.6）"
    )


@pytest.mark.asyncio
async def test_l2_export_template_none_branch_matches_legacy_builder() -> None:
    """真实调用（唯一不依赖 DB 的 handler）：`sheet=None` 的产物 == 既有构造函数的产物。

    比对 sheet 名 / 表头 / 冻结窗格，而不是 xlsx 字节 —— openpyxl 会把生成时刻写进
    `docProps/core.xml`，两次构建的字节天然不同，按字节比会得到一条恒红的假判据。
    """
    resp = await l2_host.l2_export_template(wp_id="wp-任务5.2", sheet=None, current_user=None)
    body = b"".join([chunk async for chunk in resp.body_iterator])
    assert resp.media_type == (
        "application/vnd.openxmlformats-officedocument.spreadsheetml.sheet"
    )
    got = load_workbook(io.BytesIO(body))
    buffer = io.BytesIO()
    l2_host._create_template_wb().save(buffer)
    want = load_workbook(io.BytesIO(buffer.getvalue()))

    assert got.sheetnames == want.sheetnames == [l2_host._SHEET_NAME]
    assert [c.value for c in got.active[1]] == [c.value for c in want.active[1]] == l2_host._HEADERS
    assert got.active.freeze_panes == want.active.freeze_panes == "A2"


@pytest.mark.parametrize("short", sorted(_HOSTS))
def test_delegate_unwraps_shared_impl_defaults(short: str) -> None:
    """委派时补的默认值必须是**剥掉 `Query(...)` 包装后的真值**（否则静默留幽灵行）。

    共享实现的 `import-data` 端点签名是 `strategy: str = Query("overwrite")`。直接调用
    该 callable 而不补参数，`strategy` 会拿到 `Query` 包装对象本身 ⇒ 共享实现的
    `_should_purge_residual` 比对 `== "overwrite"` 不成立 ⇒ 残留族键不清 = 幽灵行，
    而接口仍 200。这里钉住「剥出来的是 str」这一步。
    """
    module, _, _ = _HOSTS[short]
    endpoint = module._X3_SHAPE_A["import-data"]
    par = inspect.signature(endpoint).parameters["strategy"]
    unwrapped = getattr(par.default, "default", par.default)
    assert isinstance(unwrapped, str) and unwrapped, (
        f"{short}: 共享实现 import-data 的 strategy 默认值剥出 {unwrapped!r}（期望非空 str）"
    )
    assert type(par.default).__name__ == "Query", (
        f"{short}: strategy 默认值不再是 Query 包装（{type(par.default).__name__}）—— "
        "共享实现签名变了，委派点的剥包装逻辑须同步复核"
    )

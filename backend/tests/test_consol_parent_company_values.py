"""任务 13：有分公司时母公司汇总数；P13 无分公司输出逐字节不变。

真实 SQLite + 真 ORM 行，贯通三码树、审批过滤、node_measures、报表公式与 DisclosureEngine。
"""
from __future__ import annotations

import json
import uuid
from decimal import Decimal

import pytest
import pytest_asyncio
import sqlalchemy as sa
from openpyxl import Workbook
from sqlalchemy.ext.asyncio import AsyncSession, async_sessionmaker, create_async_engine

import app.models  # noqa: F401
import tests.conftest  # noqa: F401  # 注册全部模型与 SQLite 方言补丁
from app.models.audit_platform_models import AccountCategory, TrialBalance
from app.models.base import Base, ProjectStatus
from app.models.consolidation_models import EliminationEntry, EliminationEntryType, ReviewStatusEnum
from app.models.core import Project
from app.models.report_models import FinancialReport, FinancialReportType, ReportConfig
from app.services.disclosure_engine import DisclosureEngine
from app.services.parent_company_note_sections import PARENT_SOURCE_META_KEY, resolve_parent_scope_for_notes
from app.services.parent_company_scope import resolve_parent_company_context
from app.services.parent_company_values import load_parent_company_values
from app.services.report_excel_exporter import ReportExcelExporter

Y = 2025
D = Decimal
_engine = create_async_engine("sqlite+aiosqlite:///:memory:", echo=False)


@pytest_asyncio.fixture
async def db() -> AsyncSession:
    async with _engine.begin() as conn:
        await conn.run_sync(Base.metadata.drop_all)
        await conn.run_sync(Base.metadata.create_all)
    factory = async_sessionmaker(_engine, class_=AsyncSession, expire_on_commit=False)
    async with factory() as session:
        yield session


def project(code: str, name: str, scope: str, *, parent: str | None = None, relation: str | None = None) -> Project:
    return Project(
        id=uuid.uuid4(), name=f"{name}_{scope}", client_name=name, company_code=code,
        report_scope=scope, audit_year=Y, parent_company_code=parent,
        relation_to_parent=relation, ultimate_company_code=code if parent is None else "G",
        template_type="soe", status=ProjectStatus.execution,
    )

def tb(p: Project, code: str, name: str, category: AccountCategory, audited: str, unadjusted: str, opening: str) -> TrialBalance:
    return TrialBalance(
        id=uuid.uuid4(), project_id=p.id, year=Y, company_code=p.company_code or "",
        standard_account_code=code, account_name=name, account_category=category,
        audited_amount=D(audited), unadjusted_amount=D(unadjusted), opening_balance=D(opening),
    )


def entry(p: Project, number: str, amount: str, status: ReviewStatusEnum) -> EliminationEntry:
    value = D(amount)
    return EliminationEntry(
        id=uuid.uuid4(), project_id=p.id, year=Y, entry_no=number,
        entry_type=EliminationEntryType.other, account_code="1001", entry_group_id=uuid.uuid4(),
        lines=[
            {"account_code": "1001", "account_name": "货币资金", "debit_amount": str(value), "credit_amount": "0"},
            {"account_code": "2001", "account_name": "应付账款", "debit_amount": "0", "credit_amount": str(value)},
        ],
        debit_amount=value, credit_amount=value, branch_entity_code="G", review_status=status,
    )


def config() -> ReportConfig:
    return ReportConfig(
        id=uuid.uuid4(), applicable_standard="soe_consolidated",
        report_type=FinancialReportType.balance_sheet, row_number=1,
        row_code="BS-PARENT", row_name="货币资金", formula="TB('1001','期末余额')",
        formula_category="auto_calc", is_total_row=False, indent_level=0,
    )


async def seed_branch(db: AsyncSession) -> dict[str, Project]:
    rows = {
        "G": project("G", "甲集团", "consolidated"),
        "G_s": project("G", "甲集团", "standalone"),
        "GB": project("GB", "甲分公司", "standalone", parent="G", relation="branch"),
        "S": project("S", "乙子公司", "standalone", parent="G", relation="subsidiary"),
    }
    db.add_all(rows.values())
    db.add_all([
        tb(rows["G_s"], "1001", "货币资金", AccountCategory.asset, "100", "90", "80"),
        tb(rows["GB"], "1001", "货币资金", AccountCategory.asset, "20", "18", "15"),
        tb(rows["S"], "1001", "货币资金", AccountCategory.asset, "1000", "900", "800"),
        entry(rows["G"], "BR-A", "5", ReviewStatusEnum.approved),
        entry(rows["G"], "BR-D", "99", ReviewStatusEnum.draft),
        config(),
    ])
    await db.commit()
    return rows

@pytest.mark.asyncio
async def test_branch_parent_values_share_one_kernel(db: AsyncSession) -> None:
    rows = await seed_branch(db)
    context = await resolve_parent_company_context(db, rows["G"])
    assert context.has_branches is True
    assert context.parent_node_key == "G:parent"
    assert [node.node_key for node in context.parent_node.children] == ["G:branch_elim", "G:hq", "GB:branch"]

    values = await load_parent_company_values(db, rows["G"].id, Y, context)
    assert values.audited["1001"] == D("125.00"), "本部100 + 分公司20 + approved母分差额5"
    assert values.unadjusted["1001"] == D("108"), "未审数只取本部90 + 分公司18"
    assert values.opening["1001"] == D("95"), "期初数只取本部80 + 分公司15"
    assert values.audited["1001"] != D("1125"), "子公司 S 不属于母公司汇总节点"

    exporter = ReportExcelExporter(db)
    audited = await exporter._load_parent_row_index(rows["G"], Y, ["balance_sheet"], mode="audited")
    unadjusted = await exporter._load_parent_row_index(rows["G"], Y, ["balance_sheet"], mode="unadjusted")
    assert audited["BS-PARENT"]["current_period_amount"] == D("125.00")
    assert unadjusted["BS-PARENT"]["current_period_amount"] == D("108.00")
    assert audited["BS-PARENT"]["prior_period_amount"] is None
    assert audited["BS-PARENT"]["_prior_blank_reason"] == "母公司汇总上年数未结转"


@pytest.mark.asyncio
async def test_branch_parent_note_cache_and_provenance(db: AsyncSession) -> None:
    rows = await seed_branch(db)
    engine = DisclosureEngine(db)
    scope = await resolve_parent_scope_for_notes(db, rows["G"].id, Y)
    assert scope.has_branches is True and scope.parent_node_key == "G:parent"
    engine._parent_scope_cache = scope
    ctx = await engine._build_resolver_ctx(rows["G"].id, Y, "十六、应收账款")
    assert ctx["project_id"] == rows["G_s"].id, "其他 resolver 身份仍保持母公司 standalone"
    assert ctx["_tb_cache"]["1001"] == {"audited": 125.0, "unadjusted": 108.0, "opening": 95.0}
    assert ctx["_tb_cache"]["货币资金"] is ctx["_tb_cache"]["1001"]
    assert ctx[PARENT_SOURCE_META_KEY] == {
        "source_project_name": rows["G_s"].name,
        "source_company_code": "G",
        "source_scope": "parent_aggregate",
        "source_node_key": "G:parent",
        "includes_branches": True,
    }

@pytest.mark.asyncio
async def test_parent_prior_is_forced_blank_with_visible_reason(db: AsyncSession) -> None:
    rows = await seed_branch(db)
    exporter = ReportExcelExporter(db)
    index = await exporter._load_parent_row_index(rows["G"], Y, ["balance_sheet"], mode="audited")
    should_write, value = exporter._resolve_fill_value(
        index["BS-PARENT"], "prior", "BS-PARENT", {"BS-PARENT": "zero"}
    )
    assert (should_write, value) == (False, None), "fill_empty_as=zero 不得把母公司汇总上年数复活成 0"

    wb = Workbook()
    ws = wb.active
    ws["A1"] = "{{row:BS-PARENT:prior:parent}}"
    assert exporter._fill_by_placeholders(
        ws, {}, {}, True, {"BS-PARENT": "zero"}, parent_row_index=index
    ) is True
    assert ws["A1"].value is None
    assert ws["A1"].comment is not None
    assert ws["A1"].comment.text == "母公司汇总上年数未结转"


def canonical_bytes(value: object) -> bytes:
    return json.dumps(value, ensure_ascii=False, sort_keys=True, separators=(",", ":"), default=str).encode("utf-8")


@pytest.mark.asyncio
async def test_p13_no_branch_is_byte_identical_to_old_standalone_path(db: AsyncSession) -> None:
    consolidated = project("H", "无分支集团", "consolidated")
    standalone = project("H", "无分支集团", "standalone")
    db.add_all([consolidated, standalone, tb(
        standalone, "1001", "货币资金", AccountCategory.asset, "77", "66", "55"
    )])
    db.add(FinancialReport(
        id=uuid.uuid4(), project_id=standalone.id, year=Y,
        report_type=FinancialReportType.balance_sheet, row_code="BS-002", row_name="货币资金",
        current_period_amount=D("77"), prior_period_amount=D("55"),
    ))
    await db.commit()

    context = await resolve_parent_company_context(db, consolidated)
    assert context.has_branches is False and context.parent_node_key == "H:parent"

    exporter = ReportExcelExporter(db)
    old_data = await exporter._load_report_data(standalone.id, Y, ["balance_sheet"], mode="audited")
    old_index = {row["row_code"]: row for report_rows in old_data.values() for row in report_rows}
    new_index = await exporter._load_parent_row_index(consolidated, Y, ["balance_sheet"], mode="audited")
    assert canonical_bytes(new_index) == canonical_bytes(old_index)

    engine = DisclosureEngine(db)
    old_cache = await engine._parent_tb_cache(standalone.id, Y)
    scope = await resolve_parent_scope_for_notes(db, consolidated.id, Y)
    assert scope.has_branches is False
    engine._parent_scope_cache = scope
    ctx = await engine._build_resolver_ctx(consolidated.id, Y, "十六、应收账款")
    assert canonical_bytes(ctx["_tb_cache"]) == canonical_bytes(old_cache)
    assert ctx[PARENT_SOURCE_META_KEY] == {
        "source_project_name": standalone.name,
        "source_company_code": "H",
        "source_scope": "standalone",
    }


@pytest.mark.asyncio
async def test_child_has_branch_does_not_turn_root_parent_into_aggregate(db: AsyncSession) -> None:
    """根 mode=mixed 也不能作为母公司有分公司的判据。"""
    h = project("G", "根集团", "consolidated")
    hs = project("G", "根集团", "standalone")
    child = project("A", "子公司", "standalone", parent="G", relation="subsidiary")
    child_branch = project("AB", "子公司分部", "standalone", parent="A", relation="branch")
    db.add_all([h, hs, child, child_branch])
    await db.commit()

    context = await resolve_parent_company_context(db, h)
    assert context.tree.mode == "mixed"
    assert context.parent_node_key == "G:parent"
    assert context.parent_node.kind == "data"
    assert context.has_branches is False, "只能看根的 parent 节点，禁用整棵树 mode 推断"

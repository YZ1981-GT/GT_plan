# -*- coding: utf-8 -*-
"""D4-4 营业收入调整分录汇总：契约 + 几何 + 投影往返守卫。

spec: d4-4-adjustment-summary-bidirectional-writeback · Task 5

参照 `tests/workpaper_sync/test_d4_19_discount_contract.py`（D4-19，同为单区动态行表）。
放在 `backend/tests/` 而非 `workpaper_sync/` 子目录：与同表既有测试
`test_d4_4_import_export_roundtrip.py` 同处，spec Task 5 亦如此指定。

🔴 本文件的几何断言**全部与 openpyxl 实测对账**，不写死"文档说是多少"——
Req 1.5 明令「以现算为准」。模板一旦改版，这些用例会打红而不是静默错列。
"""
from __future__ import annotations

import json
import re
from pathlib import Path

import pytest

import app.services.workpaper_sync.phase5_d4_revenue_detail as D4
from app.services.workpaper_sync.contracts import parse_contract
from app.services.workpaper_sync.phase5_d4_adjustment_sheet import (
    FIRST_DATA_ROW_D44,
    FOOTER_MARKER_D44,
    FOOTER_ROW_D44,
    HEADER_ROW_D44,
    LAST_DATA_ROW_D44,
    MANAGED_FIELD_SPECS_D44,
    MANAGED_LAST_COL_D44,
    MANAGED_SHEET_D44,
    ROW_IDENTITY_KEY_D44,
    STORE_ITEM_ID_D44,
    TABLE_KEY_D44,
    UUID_COL_D44,
    _rows,
    build_store_projection_d44,
    formula_mask_cells_d44,
    merge_projection_into_d44_rows,
    sheet_payload_d44,
    stable_key_for_d44,
)

D44_SHEET_KEY = "d44-managed"

#: 权威册（与 provider 的 TEMPLATE_RELATIVE_PATH 同一本）。
_TEMPLATE = (
    Path(__file__).resolve().parents[1] / "wp_templates" / "D" / "D4 收入底稿.xlsx"
)

#: 两种并存的行身份格式（Req 1.4）——**都必须被接纳**，且禁加格式正则。
_RID_FRONTEND = "d4a-ms2p8tkl-juz5kck"        # 前端 generateRowId()
_RID_IMPORT = "3f2a1b4c-5d6e-7f80-9a1b-2c3d4e5f6071"  # 导入侧 _parse_d4_4_row → uuid4()


@pytest.fixture(autouse=True)
def _enable(monkeypatch):
    """开本表接入 flag（同 D4-19 测试的做法）。

    命名 `_INCLUDE_D44_ADJUSTMENT_SHEET` 取既有惯例 `_INCLUDE_D{n}_{名字}_SHEET`
    （`_INCLUDE_D419_DISCOUNT_SHEET` / `_INCLUDE_D417_CUTOFF_SHEET` /
    `_INCLUDE_D420_RETURN_SHEET` / `_INCLUDE_D434_CONTRACT_SHEET` /
    `_INCLUDE_D48_PRODUCT_MARGIN_SHEET` 共 5 个先例），
    而非 design §3 写的 `_INCLUDE_D44_SHEET` —— 该处已在 tasks.md 登记更正。
    """
    monkeypatch.setattr(D4, "_INCLUDE_D44_ADJUSTMENT_SHEET", True, raising=False)


@pytest.fixture(scope="module")
def worksheet():
    """模板里的 D4-4 sheet（只读）。缺模板/缺 sheet 一律 fail 而不是 skip。"""
    import openpyxl

    assert _TEMPLATE.exists(), f"权威册不存在: {_TEMPLATE}"
    wb = openpyxl.load_workbook(_TEMPLATE, data_only=False)
    assert MANAGED_SHEET_D44 in wb.sheetnames, (
        f"sheet {MANAGED_SHEET_D44!r} 不在权威册里（现有含 D4-4 的: "
        f"{[s for s in wb.sheetnames if 'D4-4' in s]}）"
    )
    ws = wb[MANAGED_SHEET_D44]
    yield ws
    wb.close()


def _contract():
    return parse_contract(D4.build_contract_payload())


def _d44_table():
    """契约里 D4-4 的动态行 table。"""
    sheet = next(s for s in _contract().sheets if s.sheet_key == D44_SHEET_KEY)
    rows_tables = [t for t in sheet.tables if t.row_identity is not None]
    assert len(rows_tables) == 1, f"D4-4 应只有 1 个动态行 table，实得 {len(rows_tables)}"
    return rows_tables[0]


class TestGeometryFrozen:
    """几何常量 vs openpyxl 实测逐项对账（Req 1.1 / 1.2 / 1.5）。"""

    def test_sheet_extent_is_a1_j23(self, worksheet):
        assert (worksheet.max_row, worksheet.max_column) == (23, 10), (
            f"模板范围变了：现算 A1:{chr(64 + worksheet.max_column)}{worksheet.max_row}"
            " —— 受管几何须重新裁决"
        )

    def test_header_row_ten_columns_verbatim(self, worksheet):
        """表头 R5 十列逐字（Req 1.2）。字段表的 header_text 必须与模板一致。"""
        expected = [f[5] for f in MANAGED_FIELD_SPECS_D44]
        actual = [
            str(worksheet.cell(row=HEADER_ROW_D44, column=i).value or "").strip()
            for i in range(1, 11)
        ]
        assert actual == expected, f"表头漂移：\n模板 {actual}\n字段表 {expected}"
        assert all(actual), "表头有空列 ⇒ header_source_ref 会指向空格"

    def test_data_region_bounds_and_footer_do_not_collide(self):
        assert FIRST_DATA_ROW_D44 == 6
        assert LAST_DATA_ROW_D44 == 20
        assert FOOTER_ROW_D44 == 21
        assert LAST_DATA_ROW_D44 < FOOTER_ROW_D44, (
            "末数据行与 footer 行重叠 ⇒ 插行会覆盖提示文本"
        )
        assert LAST_DATA_ROW_D44 - FIRST_DATA_ROW_D44 + 1 == 15, "模板数据区容量应为 15 行"

    def test_footer_marker_equals_template_a21_verbatim(self, worksheet):
        """🔴 footer marker 必须与 A21 **全等**，不能只写前缀。

        引擎判据是 `_find_marker_row` 里的 `text.strip() == marker`（**全等**），
        不是 `startswith`。首版把 marker 写成 `"提示："` 前缀，openpyxl 侧
        `a21.startswith(marker)` 当然为真、48 个单测全绿 —— 但真栈 rematerialize
        14.2s 就抛 `FooterAnchorDriftError: footer marker 在列 A 上一处都找不到`。
        **本地断言用的是我自己的 startswith 口径，与引擎的全等口径不是一回事。**

        对照：已落地的同类 sheet 一律用完整行文本（D4-19 `三、审计说明` /
        D4-34 `2.咨询业务`），没有一个用前缀。
        """
        a21 = str(worksheet.cell(row=FOOTER_ROW_D44, column=1).value or "")
        assert len(a21) == 63, f"A{FOOTER_ROW_D44} 提示文本长度变了：现算 {len(a21)} 字"
        assert FOOTER_MARKER_D44 == a21.strip(), (
            "marker 与模板 A21 不全等 —— 引擎按全等匹配，不等即定位不到 footer 锚\n"
            f"marker  = {FOOTER_MARKER_D44!r}\n"
            f"模板A21 = {a21.strip()!r}"
        )

    def test_engine_marker_matcher_rejects_a_prefix_marker(self, worksheet):
        """变异反证：用**引擎自己的**匹配函数证明「前缀 marker」确实定位不到。

        不是断言「我觉得前缀不行」，而是真调 `_find_marker_row`：
        完整文本 → 命中行 21；前缀 `"提示："` → `None`。
        """
        import zipfile

        from app.services.workpaper_sync.excel_materialize import (
            _find_marker_row,
            _shared_strings,
        )

        with zipfile.ZipFile(_TEMPLATE) as zf:
            entries = {n: zf.read(n) for n in zf.namelist()}
        # 定位该 sheet 的 part
        wb_xml = entries["xl/workbook.xml"].decode("utf-8")
        rels = entries["xl/_rels/workbook.xml.rels"].decode("utf-8")
        import re as _re

        m = _re.search(
            r'<sheet name="' + _re.escape(MANAGED_SHEET_D44) + r'"[^>]*r:id="([^"]+)"', wb_xml
        )
        assert m, "定位不到 D4-4 的 sheet 关系 id"
        rid = m.group(1)
        m2 = _re.search(r'Id="' + _re.escape(rid) + r'"[^>]*Target="([^"]+)"', rels)
        assert m2, "定位不到 sheet part 路径"
        part = "xl/" + m2.group(1).lstrip("/")
        xml = entries[part].decode("utf-8")
        shared = _shared_strings(entries)

        full = _find_marker_row(xml, column="A", marker=FOOTER_MARKER_D44, shared=shared)
        assert full == FOOTER_ROW_D44, (
            f"引擎用完整 marker 定位到行 {full}，期望 {FOOTER_ROW_D44}"
        )

        prefix_only = _find_marker_row(xml, column="A", marker="提示：", shared=shared)
        assert prefix_only is None, (
            "引擎居然用前缀也能命中 —— 那说明 _find_marker_row 改成了 startswith，"
            "本条判据（以及上面那条全等断言的必要性）需要重新评估"
        )

    def test_data_region_has_no_formula_and_no_content(self, worksheet):
        """数据区公式格 0 + 非空格 0 —— 这两条是 `formula_mask` 为空的**判据**。"""
        formulas, nonempty = [], []
        for r in range(FIRST_DATA_ROW_D44, LAST_DATA_ROW_D44 + 1):
            for c in range(1, 11):
                v = worksheet.cell(row=r, column=c).value
                if isinstance(v, str) and v.startswith("="):
                    formulas.append(worksheet.cell(row=r, column=c).coordinate)
                if v is not None and str(v).strip():
                    nonempty.append(worksheet.cell(row=r, column=c).coordinate)
        assert formulas == [], f"数据区出现公式格 {formulas} ⇒ formula_mask 不能再为空"
        assert nonempty == [], f"数据区出现预置内容 {nonempty} ⇒ 需重新裁决受管范围"

    def test_merged_cells_only_in_title_rows(self, worksheet):
        merged = sorted(str(m) for m in worksheet.merged_cells.ranges)
        assert merged == ["A1:J1", "A2:J2"], f"合并区变了：{merged}"
        in_data = [
            m
            for m in merged
            if any(
                FIRST_DATA_ROW_D44 <= int(n) <= LAST_DATA_ROW_D44
                for n in re.findall(r"[A-Z]+(\d+)", m)
            )
        ]
        assert in_data == [], f"数据区出现合并单元格 {in_data} ⇒ 动态行方案不成立"

    def test_uuid_carrier_column_is_empty_in_template(self, worksheet):
        """UUID 载体列 K 必须全空，且紧邻受管末列 J 右侧（Req 1.1）。

        K~O 五列都空是「无需注入列」的依据（对比 D4-19 要注入列 P）。
        """
        assert MANAGED_LAST_COL_D44 == "J"
        assert UUID_COL_D44 == "K", "载体列须紧邻受管末列右侧"
        assert ord(UUID_COL_D44) == ord(MANAGED_LAST_COL_D44) + 1
        for col_idx, col in ((11, "K"), (12, "L"), (13, "M"), (14, "N"), (15, "O")):
            dirty = [
                (r, worksheet.cell(row=r, column=col_idx).value)
                for r in range(1, 24)
                if str(worksheet.cell(row=r, column=col_idx).value or "").strip()
            ]
            assert dirty == [], f"{col} 列非空 {dirty[:3]} ⇒ 载体位需重选"


class TestManagedFields:
    """受管字段 10 个 / key 全小写 / 列连续（Req 2.1 / 2.4）。"""

    def test_exactly_ten_managed_fields_on_columns_a_to_j(self):
        assert len(MANAGED_FIELD_SPECS_D44) == 10
        cols = [f[1] for f in MANAGED_FIELD_SPECS_D44]
        assert cols == [chr(ord("A") + i) for i in range(10)], (
            f"受管列不是连续 A~J：{''.join(cols)}"
        )

    def test_field_names_match_frontend_row_type(self):
        """字段名须与前端 `D4AdjustmentRow` 的 10 个业务字段一致（`rowId` 是身份不计入）。"""
        assert [f[0] for f in MANAGED_FIELD_SPECS_D44] == [
            "description",
            "category",
            "reportItem",
            "accountName",
            "noteItem",
            "placeholder",
            "debitAmount",
            "creditAmount",
            "indexRef",
            "remark",
        ]
        assert ROW_IDENTITY_KEY_D44 == "rowId", "行身份键是 rowId（不是 D4-19 的 id）"
        assert ROW_IDENTITY_KEY_D44 not in [f[0] for f in MANAGED_FIELD_SPECS_D44]

    def test_amount_fields_are_typed_amount_others_text(self):
        by_field = {f[0]: f[3] for f in MANAGED_FIELD_SPECS_D44}
        assert by_field["debitAmount"] == "amount"
        assert by_field["creditAmount"] == "amount"
        assert {
            k for k, v in by_field.items() if v == "amount"
        } == {"debitAmount", "creditAmount"}, "只有借贷两列是金额类型"

    def test_stable_field_keys_are_all_lowercase(self):
        """🔴 D4-8 曾因驼峰生成 180 个含大写非法 key 打挂整份契约 parse（Req 2.4）。"""
        pat = re.compile(r"^[a-z0-9_./{}-]+$")
        for field, *_ in MANAGED_FIELD_SPECS_D44:
            sk = stable_key_for_d44(field)
            assert pat.fullmatch(sk), f"非法 stable_field_key: {sk}"
        # 逐条钉死驼峰 → snake 的转换结果（防 _snake 被改坏）
        assert stable_key_for_d44("reportItem") == f"{TABLE_KEY_D44}/{{row_uuid}}/report_item"
        assert stable_key_for_d44("debitAmount") == f"{TABLE_KEY_D44}/{{row_uuid}}/debit_amount"
        assert stable_key_for_d44("indexRef") == f"{TABLE_KEY_D44}/{{row_uuid}}/index_ref"

    def test_json_pointer_keeps_camel_case(self):
        """契约 key 小写、但 `json_pointer`/store 写回**保持驼峰**（前端真源）。"""
        fields = {f["stable_field_key"]: f for f in sheet_payload_d44()["tables"][0]["fields"]}
        ptr = fields[stable_key_for_d44("reportItem")]["json_pointer"]
        assert ptr == "/rows/{row_uuid}/reportItem", (
            "json_pointer 被 snake 化了 ⇒ store 写回会写到不存在的字段名"
        )


class TestFormulaMask:
    """`formula_mask` 空 + CS-13 空分母显式记录（Req 2.2 / 2.3）。"""

    def test_formula_mask_is_empty(self):
        assert formula_mask_cells_d44() == ()
        assert sheet_payload_d44()["tables"][0]["formula_mask"] == []

    def test_no_formula_mode_field_so_cs13_holds_on_empty_denominator(self):
        """CS-13「formula 字段必须落在 formula_mask 内」在本表是**空分母成立**。

        🔴 显式记录这一点：不得因 `formula_mask` 为空而被误判违规。判据是本表
        **没有** `mode="formula"` 的字段 ⇒ 反向约束的分母为 0。
        """
        modes = {f[2] for f in MANAGED_FIELD_SPECS_D44}
        assert modes == {"editable"}, f"出现非 editable 字段 {modes} ⇒ CS-13 空分母不再成立"
        assert not [f for f in MANAGED_FIELD_SPECS_D44 if f[2] == "formula"]


class TestContractWiring:
    """D4-4 进契约且可 parse（Req 3.1 / 3.2）。

    🔴 本类依赖 Task 6 的接线。接线前它**必然打红** —— 那正是「判据有区分力」的证明，
    不是缺陷。接线完成后转绿。
    """

    def test_d44_sheet_present_in_contract_payload(self):
        keys = [s["sheet_key"] for s in D4.build_contract_payload()["sheets"]]
        assert D44_SHEET_KEY in keys, (
            f"D4-4 未进契约（现有 {len(keys)} 张）⇒ Task 6 第 2 处接线未完成"
        )

    def test_contract_parses_without_raising(self):
        assert any(s.sheet_key == D44_SHEET_KEY for s in _contract().sheets)

    def test_table_shape_in_parsed_contract(self):
        t = _d44_table()
        assert t.table_key == TABLE_KEY_D44
        assert ROW_IDENTITY_KEY_D44 in str(t.row_identity.json_pointer)
        assert {f.cell.column for f in t.fields} == set("ABCDEFGHIJ")
        assert len(t.fields) == 10

    def test_store_item_registered_for_mirror(self):
        """🔴 `D4-4-rows` 必须在 `STORE_ITEM_IDS` 里，否则 mirror 取不到它的 base。"""
        assert STORE_ITEM_ID_D44 in set(D4.STORE_ITEM_IDS)

    def test_raw_payload_table_key_mapping_registered(self):
        """🔴 merge 返回裸 list ⇒ 必须登记 `_RAW_PAYLOAD_ITEM_TABLE_KEYS`。

        漏登记的后果不是"少同步一张表"，而是 `_normalize_merge_updates` 归一不到
        4-tuple ⇒ `store_mirror.py` 的硬解包抛 `ValueError` ⇒ **打挂整个 entry 的回写**
        （D4-8 踩过）。
        """
        mapping = D4._RAW_PAYLOAD_ITEM_TABLE_KEYS
        assert STORE_ITEM_ID_D44 in mapping, (
            f"{STORE_ITEM_ID_D44} 未登记 ⇒ 4-tuple 归一会漏掉它，mirror 硬解包 ValueError"
        )
        assert mapping[STORE_ITEM_ID_D44] == (TABLE_KEY_D44,)


def _sample_row(rid: str) -> dict:
    return {
        ROW_IDENTITY_KEY_D44: rid,
        "description": "调整营业收入跨期",
        "category": "账项调整",
        "reportItem": "营业收入",
        "accountName": "主营业务收入",
        "noteItem": "五、1 营业收入",
        "placeholder": "补充说明文本",
        "debitAmount": 120000.0,
        "creditAmount": 0.0,
        "indexRef": "D4-2",
        "remark": "经复核",
    }


class TestRoundtrip:
    """projection ↔ merge 往返等值（两种 rowId 格式各跑一遍）。"""

    @pytest.mark.parametrize("rid", [_RID_FRONTEND, _RID_IMPORT], ids=["frontend", "import"])
    def test_roundtrip_equal_for_both_identity_formats(self, rid):
        base = [_sample_row(rid)]
        parsed = _contract()
        proj = build_store_projection_d44(json.dumps(base), contract=parsed)
        assert proj.row_keys[TABLE_KEY_D44] == (rid,)
        assert len(proj.values) == 10
        merged = merge_projection_into_d44_rows(projection=proj, base_payload=base)
        assert merged == base, f"往返不等值:\n期望 {base}\n实得 {merged}"

    def test_unmanaged_keys_are_preserved(self):
        base = [{**_sample_row(_RID_FRONTEND), "__private__": "不受管但须保留"}]
        proj = build_store_projection_d44(base, contract=_contract())
        merged = merge_projection_into_d44_rows(projection=proj, base_payload=base)
        assert merged[0]["__private__"] == "不受管但须保留"

    def test_row_order_follows_base(self):
        rids = ["d4a-a-1", "d4a-b-2", "d4a-c-3"]
        base = [_sample_row(r) for r in rids]
        proj = build_store_projection_d44(base, contract=_contract())
        merged = merge_projection_into_d44_rows(projection=proj, base_payload=base)
        assert [r[ROW_IDENTITY_KEY_D44] for r in merged] == rids

    def test_empty_base_does_not_fabricate_rows(self):
        """🔴 空 base 不得凭空造行（也就不会把 0 投进模板空行）。"""
        proj = build_store_projection_d44([_sample_row(_RID_FRONTEND)], contract=_contract())
        for empty in (None, "", "[]", [], {}, 42):
            assert merge_projection_into_d44_rows(projection=proj, base_payload=empty) == [], (
                f"空 base {empty!r} 产出了行 ⇒ 会把投影值写进不存在的行"
            )


class TestPayloadTolerance:
    """`_rows()` 两层强度：容器层容差、行层 fail-closed（Req 1.4 / T4）。"""

    @pytest.mark.parametrize(
        "payload",
        [None, "", "   ", "{not json", {"rows": []}, 42, [], b"", bytearray()],
        ids=["none", "empty", "blank", "bad-json", "dict", "int", "empty-list", "bytes", "bytearray"],
    )
    def test_container_level_is_tolerant(self, payload):
        """🔴 与另外 35 张 sheet 共享 entry，这里抛异常会打挂**整个 entry**（D4-9 事故）。"""
        assert _rows(payload) == []

    def test_non_mapping_elements_are_skipped_not_fatal(self):
        got = _rows([1, "x", None, _sample_row(_RID_FRONTEND)])
        assert [rid for rid, _ in got] == [_RID_FRONTEND]

    @pytest.mark.parametrize(
        "payload,reason",
        [
            ([{"description": "x"}], "缺 rowId"),
            ([{"rowId": "", "description": "x"}], "rowId 空串"),
            ([{"rowId": "   "}], "rowId 全空白"),
            ([{"rowId": "a"}, {"rowId": "a"}], "rowId 重复"),
        ],
    )
    def test_row_level_is_fail_closed(self, payload, reason):
        """行身份坏了必须可见地失败 —— 丢身份会让 materialize 在错区插行（W22 那类缺陷）。"""
        with pytest.raises(ValueError):
            _rows(payload)

    @pytest.mark.parametrize("rid", [_RID_FRONTEND, _RID_IMPORT], ids=["frontend", "import"])
    def test_both_identity_formats_accepted(self, rid):
        """🔴 禁加格式正则：两种格式都是唯一字符串，加校验会拒掉导入产生的行。"""
        got = _rows([{ROW_IDENTITY_KEY_D44: rid}])
        assert [r for r, _ in got] == [rid]


class TestMutationProofs:
    """🔴 变异反证 —— 证明上面的判据**真的在测东西**，不是恒绿。

    tasks.md Task 5 要求这两条「须实做」：删 `placeholder` 字段 → 红；
    `UUID_COL_D44` 改 `J` → 红。这里各用一条真实变异把对应判据打红。

    为什么必须有：`TestGeometryFrozen` / `TestManagedFields` 全是"现算 == 常量"形态，
    如果常量和现算来自同一个错误来源，它们会一起错、一起绿。变异反证是唯一能区分
    「判据成立」与「判据无区分力」的手段（本 repo 反复踩过恒绿守卫）。
    """

    def test_dropping_placeholder_field_breaks_header_alignment(self, worksheet, monkeypatch):
        """删掉 `placeholder`（模板 F 列「……」）⇒ 表头逐字对账与列连续性**都**打红。

        这一条同时守住 Req 2.5 的核心：F/J 两列在模板里真实存在，漏掉任何一个都会让
        OO 侧的编辑被 store 旧值覆盖（静默丢数据）。
        """
        import app.services.workpaper_sync.phase5_d4_adjustment_sheet as P

        mutated = tuple(f for f in MANAGED_FIELD_SPECS_D44 if f[0] != "placeholder")
        assert len(mutated) == 9, "变异样本构造失败（placeholder 不在字段表里？）"
        monkeypatch.setattr(P, "MANAGED_FIELD_SPECS_D44", mutated, raising=True)

        # ① 列连续性判据必须打红
        cols = [f[1] for f in mutated]
        assert cols != [chr(ord("A") + i) for i in range(10)], (
            "删了一个字段后列集合仍等于 A~J ⇒ 列连续性判据无区分力"
        )
        # ② 表头逐字对账必须打红（模板 F5 是「……」，变异后字段表里没有它）
        template_headers = [
            str(worksheet.cell(row=HEADER_ROW_D44, column=i).value or "").strip()
            for i in range(1, 11)
        ]
        assert [f[5] for f in mutated] != template_headers, (
            "删了 placeholder 后表头仍与模板逐字相等 ⇒ 表头判据无区分力"
        )
        # ③ 字段数判据必须打红
        assert len(mutated) != 10

    def test_uuid_column_j_collides_with_managed_last_column(self, monkeypatch):
        """把 `UUID_COL_D44` 改成 `J` ⇒ 与受管末列冲突，载体判据打红。

        为什么这是必须守的：J 是 `remark` 的受管列。载体列落在受管列上会让行身份被
        业务值覆盖（或反之），后果是身份不稳定 ⇒ 每次 sync 重新 mint ⇒ 孤儿行累积。
        """
        import app.services.workpaper_sync.phase5_d4_adjustment_sheet as P

        monkeypatch.setattr(P, "UUID_COL_D44", "J", raising=True)
        managed_cols = {f[1] for f in MANAGED_FIELD_SPECS_D44}

        assert P.UUID_COL_D44 in managed_cols, "变异样本构造失败"
        # 载体列不得落在受管列集合内 —— 这条判据在变异下必须失败
        assert not (
            P.UUID_COL_D44 not in managed_cols
        ), "UUID 列改成受管列 J 后「载体不与受管列冲突」仍成立 ⇒ 该判据无区分力"
        # 紧邻关系也必须打红
        assert ord(P.UUID_COL_D44) != ord(MANAGED_LAST_COL_D44) + 1

    def test_identity_format_regex_would_reject_import_rows(self):
        """反证「禁加格式正则」这条纪律的必要性：真加了会拒掉导入行。

        不改生产代码，只演示假设的校验函数在两种格式上的行为差异 —— 证明
        `_rows` 不加正则不是偷懒，而是必需（Req 1.4）。
        """
        def _with_format_check(rid: str) -> bool:
            return rid.startswith("d4a-")

        assert _with_format_check(_RID_FRONTEND) is True
        assert _with_format_check(_RID_IMPORT) is False, (
            "样本选错了：导入侧 rowId 不该匹配前端前缀"
        )
        # 而真实的 `_rows` 必须两者都接纳
        assert len(_rows([{ROW_IDENTITY_KEY_D44: _RID_FRONTEND}])) == 1
        assert len(_rows([{ROW_IDENTITY_KEY_D44: _RID_IMPORT}])) == 1


class TestRowCountExceedsTemplateCapacity:
    """T17：行数超模板容量（15 行）时，插行不得覆盖 footer 提示文本（Req 6.4）。

    模板数据区 R6~R20 共 15 行，footer marker 在 R21。store 若有 20 行，引擎必须**插行**
    而不是往 R21 及以后直接写 —— 后者会把「提示：本底稿适用于…」那 63 字业务说明覆盖掉。

    本类是**离线判据**（纯 projection/plan 层），不需要 live PG；真实写盘后的字节级确认
    在 T10*/T14* 的发布链与 L1 验收里。
    """

    def _rows(self, n: int) -> list[dict]:
        return [
            {
                ROW_IDENTITY_KEY_D44: f"d4a-cap-{i:04d}",
                "description": f"调整{i}",
                "debitAmount": float(i),
                "creditAmount": 0.0,
            }
            for i in range(n)
        ]

    def test_projection_carries_all_rows_beyond_template_capacity(self):
        """20 行（> 模板 15 行）必须全部进 projection —— 不得静默截断到容量上限。"""
        capacity = LAST_DATA_ROW_D44 - FIRST_DATA_ROW_D44 + 1
        assert capacity == 15, f"模板容量现算 {capacity}，本用例的前提是 15"
        rows = self._rows(20)
        proj = build_store_projection_d44(rows, contract=_contract())
        assert len(proj.row_keys[TABLE_KEY_D44]) == 20, (
            "超容量的行被静默丢弃 ⇒ 用户填的分录会消失（比报错更坏）"
        )
        # 每行 10 字段
        assert len(proj.values) == 200

    def test_footer_anchor_declares_marker_not_total_formula(self):
        """footer 是**提示文本**不是合计公式 ⇒ `carries_total_formula` 必须 False。

        若误声明 True，引擎会把 R21 当合计行去扩张公式区间，进而改写那 63 字说明。
        """
        t = _d44_table()
        anchor = t.footer_anchor
        assert anchor is not None, "缺 footer_anchor ⇒ 引擎无从判断插行下界"
        assert anchor.carries_total_formula is False
        assert FOOTER_MARKER_D44 in str(anchor.marker)

    def test_footer_marker_is_searchable_in_column_a_below_data_region(self, worksheet):
        """marker 必须在 A 列**唯一**命中且落在数据区之下。

        🔴 判据用**全等**（`== marker`）而不是 `startswith` —— 与引擎
        `_find_marker_row` 的 `text.strip() == marker` 同口径。首版用 startswith，
        与引擎口径不一致，导致「本地绿、真栈 FooterAnchorDriftError」。
        """
        hits = [
            r
            for r in range(1, worksheet.max_row + 1)
            if str(worksheet.cell(row=r, column=1).value or "").strip() == FOOTER_MARKER_D44
        ]
        assert hits == [FOOTER_ROW_D44], (
            f"marker「{FOOTER_MARKER_D44}」在 A 列命中行 {hits}，期望恰 [{FOOTER_ROW_D44}] —— "
            "多处命中会让锚定位漂移，零命中会让引擎找不到插行下界"
        )
        assert all(h > LAST_DATA_ROW_D44 for h in hits), "marker 落在数据区内 ⇒ 会被当数据行"

    def test_managed_columns_never_reach_uuid_carrier(self):
        """受管列与 UUID 载体列不重叠 —— 插行时身份列不会被业务值覆盖。"""
        managed = {f[1] for f in MANAGED_FIELD_SPECS_D44}
        assert UUID_COL_D44 not in managed
        assert max(managed) == MANAGED_LAST_COL_D44

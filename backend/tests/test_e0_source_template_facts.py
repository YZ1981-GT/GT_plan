"""E0 函证源模板事实守卫（openpyxl 直读源 xlsx，三向比对 + 反向自检）。

锁死三件事：

1. **E0-5 归属裁决** —— 工作簿里有两张 tab 以 `E0-5` 结尾
   （`应付银行承兑汇票发函记录表E0-5` / `银行函证其他信息核对表E0-5`），
   而 `底稿目录` 的索引号列（F 列）**只把 E0-5 给了前者**，后者在底稿目录里没有索引号。
   故凡「按编码定位 E0-5」的地方都必须落到应付银行承兑汇票发函记录表。

2. **wp_account_mapping.json 的 E0 名称** 与源模板 `底稿目录` 逐字一致
   （历史上是 D0 口径的错名：跟函控制 / 差异调节 / 替代程序）。

3. **四张发函记录表的取数口径** —— 列头逐字 + `函证结果汇总表E0-1!F8` 嵌套
   SUMIF/SUMIFS 的匹配键与取值列；前端 `importE0ListsToSummary.ts` 的
   `E0_LIST_SPECS` 必须与之对齐（跨前后端交叉锁死）。
"""
from __future__ import annotations

import json
import re
from pathlib import Path

import openpyxl
import pytest

_BACKEND = Path(__file__).resolve().parents[1]
_TEMPLATE = _BACKEND / "wp_templates" / "E" / "E0 货币资金 - 函证（Leap应对措施-函证）.xlsx"
_MAPPING = _BACKEND / "data" / "wp_account_mapping.json"
_FRONTEND_SPEC = (
    _BACKEND.parent
    / "audit-platform"
    / "frontend"
    / "src"
    / "components"
    / "workpaper"
    / "confirmation"
    / "coordination"
    / "importE0ListsToSummary.ts"
)

# 源模板 `底稿目录` D/E/F 三列（序号 / 内容 / 索引号），行 3~11
EXPECTED_INDEX = {
    "E0A": "函证程序表",
    "E0-1": "函证结果汇总表",
    "E0-2": "核实被函证单位信息",
    "E0-3": "货币资金发函记录表",
    "E0-4": "借款发函记录表",
    "E0-5": "应付银行承兑汇票发函记录表",
    "E0-6": "理财产品发函记录表",
    "E0-7": "跟函函证过程控制",
    "E0-8": "函证程序舞弊风险评价表",
}

# 发函记录表列头（逐字，第 5 行）
EXPECTED_HEADERS = {
    "货币资金发函记录表E0-3": [
        "所属科目", "索引号", "报表截止日", "开户银行", "是否函证", "账户名称",
        "银行账号", "币种", "利率(%)", "账户类型", "账户余额（原币）",
        "是否属于资金归集（资金池或其他资金管理）账户", "起始日期", "终止日期",
        "是否存在冻结、担保或其他使用限制（如是，请注明）", "备注",
    ],
    "借款发函记录表E0-4": [
        "所属科目", "索引号", "报表截止日", "开户银行", "是否函证", "借款人名称",
        "借款账号", "币种", "余额", "借款日期", "到期日期", "利率(%)",
        "抵(质)押品/担保人", "备注", "借款类型", "期末应付利息",
    ],
    "应付银行承兑汇票发函记录表E0-5": [
        "索引号", "报表截止日", "开户银行", "银行承兑汇票号码", "结算账户账号",
        "币种", "票面金额", "出票日", "到期日", "抵（质）押品",
    ],
    "理财产品发函记录表E0-6": [
        "索引号", "报表截止日", "开户行名称及收件人", "产品名称",
        "产品类型（封闭式/开放式）", "币种", "持有份额", "产品净值",
        "购买日", "到期日", "是否被用于担保或存在其他使用限制",
    ],
}


@pytest.fixture(scope="module")
def wb():
    if not _TEMPLATE.exists():  # pragma: no cover - 模板缺失时跳过而非误红
        pytest.skip(f"源模板不存在: {_TEMPLATE}")
    return openpyxl.load_workbook(_TEMPLATE, data_only=False)


@pytest.fixture(scope="module")
def catalog(wb) -> dict[str, str]:
    """{索引号: 内容} —— 底稿目录 F 列 → E 列。"""
    ws = wb["底稿目录"]
    out: dict[str, str] = {}
    for r in range(1, ws.max_row + 1):
        code = ws.cell(r, 6).value  # F 列 索引号
        name = ws.cell(r, 5).value  # E 列 内容
        if isinstance(code, str) and code.strip().startswith("E0") and isinstance(name, str):
            out[code.strip()] = name.strip()
    return out


@pytest.fixture(scope="module")
def mapping() -> dict[str, dict]:
    with open(_MAPPING, encoding="utf-8-sig") as f:
        data = json.load(f)
    return {
        e["wp_code"]: e
        for e in data.get("mappings", [])
        if str(e.get("wp_code", "")).startswith("E0")
    }


# ── 1. E0-5 归属裁决 ────────────────────────────────────────────────────────


def test_two_tabs_end_with_e0_5(wb):
    """工作簿确实存在「一码两表」：两张 tab 都以 E0-5 结尾（反向自检本守卫的前提）。"""
    tabs = [s for s in wb.sheetnames if s.endswith("E0-5")]
    assert sorted(tabs) == ["应付银行承兑汇票发函记录表E0-5", "银行函证其他信息核对表E0-5"]


def test_catalog_assigns_e0_5_to_bank_acceptance_list(catalog):
    """底稿目录把 E0-5 唯一给了应付银行承兑汇票发函记录表。"""
    assert catalog["E0-5"] == "应付银行承兑汇票发函记录表"


def test_bank_info_check_sheet_has_no_catalog_index(catalog):
    """银行函证其他信息核对表在底稿目录里没有索引号（它不是 E0-5，也不是 E0-6）。"""
    assert "银行函证其他信息核对表" not in catalog.values()
    # 其自身 T3 写 `=底稿目录!F9`，而 F9 = E0-6 = 理财产品发函记录表（源模板贴错）
    assert catalog["E0-6"] == "理财产品发函记录表"


def test_bank_info_check_sheet_index_cell_is_wrong(wb):
    """钉死源模板缺陷：核对表 T3 引 底稿目录!F9（E0-6），与其标题 E0-5 自相矛盾。"""
    ws = wb["银行函证其他信息核对表E0-5"]
    assert ws["A2"].value == "银行函证其他信息核对表"
    assert ws["T3"].value == "=底稿目录!F9"


def test_bank_info_check_sheet_has_13_clause_items(wb):
    """核对表 F6:R6 是 13 个询证函标准条款要项（逐字，含序号前缀）。"""
    ws = wb["银行函证其他信息核对表E0-5"]
    items = [ws.cell(6, c).value for c in range(6, 19)]
    assert items == [
        "3.注销账户",
        "4.本公司作为委托人的委托贷款",
        "5.本公司作为借款人的委托贷款",
        "6.（1）对外担保",
        "6.（2）接受担保",
        "8.贴现商业汇票",
        "9.托收商业汇票",
        "10.信用证",
        "11.外汇买卖合约",
        "12.托管证券或其他产权文件",
        "13.银行理财产品",
        "14.其他",
        "15.附表(资金归集)",
    ]
    # 缺 1/2/7 —— 因为 1 银行存款/2 银行借款/7 应付银行承兑汇票 由四张发函记录表承载
    assert all("1.银行存款" != i and "7." not in i for i in items)


# ── 2. wp_account_mapping 与底稿目录逐字一致 ────────────────────────────────


# 允许与底稿目录不同字面的既有命名（须逐条写明理由，防静默增长）：
# E0-2 —— E0 的被函证单位就是银行，「核实被函证银行」比模板通用名更具体，
#         且不是 D0 口径错名（D0-2 是「核实被函证单位」），保留。
_NAME_ALLOWLIST = {"E0-2": "核实被函证银行"}


@pytest.mark.parametrize("code", ["E0-1", "E0-2", "E0-3", "E0-4", "E0-5"])
def test_mapping_wp_name_matches_catalog(code, catalog, mapping):
    """已登记的 E0 子底稿 wp_name 必须与源模板底稿目录一致（曾是 D0 口径错名）。"""
    assert code in mapping, f"{code} 未在 wp_account_mapping 登记"
    expected = catalog[code]
    actual = mapping[code]["wp_name"]
    if code in _NAME_ALLOWLIST:
        assert actual == _NAME_ALLOWLIST[code], (
            f"{code} 已偏离豁免值 {_NAME_ALLOWLIST[code]!r}（实为 {actual!r}）——"
            "要么改回，要么按底稿目录 {expected!r} 命名并移出豁免"
        )
        return
    # E0-1 沿用去「表」字的历史简称，只要求是底稿目录名的前缀式简称
    assert actual and (actual == expected or expected.startswith(actual)), (
        f"{code} wp_name={actual!r} 与底稿目录 {expected!r} 不符"
    )


def test_name_allowlist_is_not_a_dumping_ground():
    """反向自检：豁免名单只许 E0-2 一条，新增必须改守卫（迫使写理由）。"""
    assert sorted(_NAME_ALLOWLIST) == ["E0-2"]


def test_mapping_has_no_d0_style_names(mapping):
    """反向自检：E0 映射里不得再出现 D0 口径的三个错名。"""
    bad = {"跟函控制", "差异调节", "替代程序"}
    hit = {c: e["wp_name"] for c, e in mapping.items() if e.get("wp_name") in bad}
    assert not hit, f"E0 映射残留 D0 口径错名: {hit}"


def test_known_missing_mapping_codes(mapping):
    """已知欠账：底稿目录有、映射缺的四条（补齐时须同步补 wp_code_overrides）。

    补齐后本断言会红 —— 那是提醒更新守卫，不是回归。
    """
    missing = sorted(set(EXPECTED_INDEX) - set(mapping))
    assert missing == ["E0-6", "E0-7", "E0-8", "E0A"], (
        f"E0 缺失映射集合已变化: {missing}（spec e0-confirmation-completion）"
    )


# ── 3. 发函记录表列头 + E0-1 聚合口径 ───────────────────────────────────────


@pytest.mark.parametrize("sheet,headers", sorted(EXPECTED_HEADERS.items()))
def test_list_sheet_headers(sheet, headers, wb):
    ws = wb[sheet]
    actual = [ws.cell(5, c).value for c in range(1, len(headers) + 1)]
    assert actual == headers


def test_e0_5_and_e0_6_have_no_confirm_flag_column(wb):
    """E0-5 / E0-6 无「是否函证」列（整表即发函记录）；E0-3 / E0-4 有 —— 反向自检。"""
    def headers(sheet: str) -> set[str]:
        ws = wb[sheet]
        return {
            str(ws.cell(5, c).value).strip()
            for c in range(1, ws.max_column + 1)
            if ws.cell(5, c).value
        }

    assert "是否函证" in headers("货币资金发函记录表E0-3")
    assert "是否函证" in headers("借款发函记录表E0-4")
    assert "是否函证" not in headers("应付银行承兑汇票发函记录表E0-5")
    assert "是否函证" not in headers("理财产品发函记录表E0-6")


def test_summary_f_column_aggregation_semantics(wb):
    """`函证结果汇总表E0-1!F8` 四条分支的匹配键与取值列（发函金额口径唯一真源）。"""
    f8 = wb["函证结果汇总表E0-1"]["F8"].value
    assert isinstance(f8, str) and f8.startswith("=IF(")

    # 银行存款/其他货币资金：SUMIF(E0-3!$G:$G(银行账号), E0-1!$E, E0-3!$K:$K(账户余额（原币）))
    assert "SUMIF('货币资金发函记录表E0-3'!$G:$G,'函证结果汇总表E0-1'!$E8,'货币资金发函记录表E0-3'!$K:$K)" in f8
    # 短/长期借款：SUMIFS(E0-4!$I:$I(余额), E0-4!$A:$A(所属科目)=E0-1!$D, E0-4!$G:$G(借款账号)=E0-1!$E)
    assert "SUMIFS('借款发函记录表E0-4'!$I:$I,'借款发函记录表E0-4'!$A:$A,'函证结果汇总表E0-1'!$D8,'借款发函记录表E0-4'!$G:$G,'函证结果汇总表E0-1'!$E8)" in f8
    # 应付票据：单条件按索引号汇总票面金额（一函多票）
    assert "SUMIF('应付银行承兑汇票发函记录表E0-5'!$A:$A,'函证结果汇总表E0-1'!$B8,'应付银行承兑汇票发函记录表E0-5'!$G:$G)" in f8
    # 理财产品：SUMIFS(产品净值, 索引号, 产品名称)
    assert "SUMIFS('理财产品发函记录表E0-6'!$H:$H,'理财产品发函记录表E0-6'!$A:$A,'函证结果汇总表E0-1'!$B8,'理财产品发函记录表E0-6'!$D:$D,'函证结果汇总表E0-1'!$E8)" in f8


def test_e0_5_branch_does_not_use_column_e(wb):
    """反向自检：应付票据分支只用 E0-1!$B（索引号），不用 $E（账号/理财产品名称）。"""
    f8 = wb["函证结果汇总表E0-1"]["F8"].value
    m = re.search(r"IF\(D8=\"应付票据\",(.*?)(?:,IF\(D8=\"理财产品\")", f8)
    assert m, "未定位到应付票据分支"
    assert "$B8" in m.group(1)
    assert "$E8" not in m.group(1)


# ── 4. 前端 E0_LIST_SPECS 交叉锁死 ─────────────────────────────────────────


def _frontend_source() -> str:
    if not _FRONTEND_SPEC.exists():  # pragma: no cover
        pytest.skip(f"前端文件不存在: {_FRONTEND_SPEC}")
    return _FRONTEND_SPEC.read_text(encoding="utf-8")


def test_frontend_specs_declare_source_sheet_names():
    src = _frontend_source()
    for sheet in EXPECTED_HEADERS:
        assert f"'{sheet}'" in src, f"前端 E0_LIST_SPECS 缺 sheetName {sheet}"


def test_frontend_specs_amount_columns_match_source():
    """金额列取源模板逐字列名（历史缺 票面金额/产品净值 → 两品种发函金额恒 0）。"""
    src = _frontend_source()
    for col in ["账户余额（原币）", "余额", "票面金额", "产品净值"]:
        assert f"'{col}'" in src, f"前端 E0_LIST_SPECS 缺金额列 {col}"


def test_frontend_specs_group_by_matches_summary_formula():
    """E0-5 单条件按索引号、E0-6 索引号+产品名（对齐 E0-1!F8）。"""
    src = _frontend_source()
    block = re.search(r"'E0-5':\s*\{(.*?)\n  \},", src, re.S)
    assert block, "未定位 E0-5 spec 块"
    assert "groupBy: 'index'" in block.group(1)
    assert "hasConfirmFlag: false" in block.group(1)
    block6 = re.search(r"'E0-6':\s*\{(.*?)\n  \},", src, re.S)
    assert block6, "未定位 E0-6 spec 块"
    assert "groupBy: 'index+accountNo'" in block6.group(1)
    assert "hasConfirmFlag: false" in block6.group(1)


def test_frontend_e0_4_subtype_prefers_account_subject():
    """E0-4 品种权威列是「所属科目」（E0-1 SUMIFS 匹配列），「借款类型」仅兜底。"""
    src = _frontend_source()
    block = re.search(r"'E0-4':\s*\{(.*?)\n  \},", src, re.S)
    assert block
    assert "subtypeKeys: ['所属科目', '借款类型']" in block.group(1)


# ── 5. E0-1 列集（30 列两级表头）──────────────────────────────────────────────

# R5 = 分组表头（只有部分列有分组名，其余为独立列）
# R6 = 字段名（每列都有）
# 总共 30 列 A(1)..AD(30)
E0_1_GROUPS = {
    # col_start: (group_name, span)
    3: ("发函询证纪要", 7),     # C..I (cols 3-9)
    10: ("1、发函信息", 5),     # J..N (cols 10-14)
    15: ("2、收到回函", 7),     # O..U (cols 15-21)
    22: ("3、回函金额确认", 4), # V..Y (cols 22-25)
}

E0_1_COLUMNS_R6 = [
    # col 1~2: 独立列（A/B）
    None,  # col 1 序号在 R5
    None,  # col 2 询证函索引号在 R5
    # 发函询证纪要组 cols 3-9
    "被询证单位名称", "账户/交易", "账号/理财产品名称",
    "发函金额（原币）", "币种", "汇率", "发函金额（本位币）",
    # 1、发函信息组 cols 10-14
    "函证方式", "发函日期", "发函单号", "收件地址", "地址核查是否一致",
    # 2、收到回函组 cols 15-21
    "是否收到回函", "回函方式", "是否相符", "回函日期", "回函快递单号",
    "回函发出地址", "发函地址与回函地址是否一致",
    # 3、回函金额确认组 cols 22-25
    "回函金额", "差异", "可确认金额（原币）", "可确认金额（本位币）",
    # 尾部独立列 cols 26-30
    None, None, None, None, None,  # R5 有值，R6 无值
]

E0_1_INDEPENDENT_COLS_R5 = {
    1: "序号",
    2: "询证函索引号",
    26: "抵押质押等事项回函说明",
    27: "其他函证事项回函是否相符",
    28: "函证不符事项说明",
    29: "不符事项检查索引号",
    30: "审计结论",
}


def test_e0_1_has_exactly_30_columns(wb):
    """E0-1 主表列数必须是 30（A..AD）。"""
    ws = wb["函证结果汇总表E0-1"]
    # R5 + R6 联合判定：至少一行有值即算有效列
    count = 0
    for c in range(1, 40):
        if ws.cell(5, c).value or ws.cell(6, c).value:
            count = c
    assert count == 30, f"E0-1 列数={count}，期望 30"


def test_e0_1_group_headers(wb):
    """R5 分组表头逐字比对。"""
    ws = wb["函证结果汇总表E0-1"]
    for col_start, (name, _span) in E0_1_GROUPS.items():
        actual = ws.cell(5, col_start).value
        assert actual == name, f"col {col_start} R5: {actual!r} != {name!r}"


def test_e0_1_independent_columns_in_r5(wb):
    """R5 独立列（不属于任何分组）逐字比对。"""
    ws = wb["函证结果汇总表E0-1"]
    for col, name in E0_1_INDEPENDENT_COLS_R5.items():
        actual = ws.cell(5, col).value
        assert actual == name, f"col {col} R5: {actual!r} != {name!r}"


def test_e0_1_field_names_in_r6(wb):
    """R6 字段名逐字比对（独立列在 R5，R6 为空）。"""
    ws = wb["函证结果汇总表E0-1"]
    for i, expected in enumerate(E0_1_COLUMNS_R6):
        col = i + 1
        actual = ws.cell(6, col).value
        if expected is None:
            # 独立列在 R5 有值，R6 应为空
            continue
        assert actual == expected, f"col {col} R6: {actual!r} != {expected!r}"


def test_e0_1_no_alternative_procedure_columns(wb):
    """E0-1 无替代程序四列（源模板没有，替代程序在「回函情况汇编」）。"""
    ws = wb["函证结果汇总表E0-1"]
    all_headers = set()
    for c in range(1, 35):
        for r in (5, 6):
            v = ws.cell(r, c).value
            if v:
                all_headers.add(str(v).strip())
    forbidden = {"是否采取替代程序", "替代程序确认金额", "替代后不可确认金额", "替代程序索引"}
    hit = all_headers & forbidden
    assert not hit, f"E0-1 不应有替代程序列: {hit}"


# ── 6. 品种矩阵（6 品种 × 6 指标）───────────────────────────────────────────

E0_MATRIX_CATEGORIES = ["银行存款", "其他货币资金", "短期借款", "长期借款", "应付票据", "理财产品"]
E0_MATRIX_METRICS = [
    "本期（期末）账面金额：",
    "抽取样本的发函金额：",
    "发函金额占账面金额的比例(%)：",
    "回函确认金额：",
    "回函可确认金额占发函金额的比例(%)：",
    "回函可确认金额占账面金额的比例(%)：",
]


def test_matrix_categories(wb):
    """品种矩阵 R28 E:J 逐字。"""
    ws = wb["函证结果汇总表E0-1"]
    actual = [ws.cell(28, c).value for c in range(5, 11)]  # E=5..J=10
    assert actual == E0_MATRIX_CATEGORIES


def test_matrix_metrics(wb):
    """品种矩阵 R29:R34 C 列指标逐字。"""
    ws = wb["函证结果汇总表E0-1"]
    actual = [ws.cell(r, 3).value for r in range(29, 35)]
    assert actual == E0_MATRIX_METRICS


def test_matrix_book_amount_row_has_no_formula(wb):
    """R29「本期（期末）账面金额」各品种列无公式 = 手工填写（理财产品无固定科目，宁缺勿造）。"""
    ws = wb["函证结果汇总表E0-1"]
    for c in range(5, 11):
        v = ws.cell(29, c).value
        assert v is None or not str(v).startswith("="), (
            f"R29 col {c} 有公式 {v!r}，但源模板该行应为手工填写"
        )


def test_matrix_sample_amount_uses_sumif_on_column_f(wb):
    """R30「抽取样本的发函金额」应为 SUMIF(D:D, 品种, F:F) 形态。"""
    ws = wb["函证结果汇总表E0-1"]
    for c in range(5, 11):
        v = ws.cell(30, c).value
        assert v and str(v).startswith("=") and "SUMIF" in str(v), (
            f"R30 col {c} 应有 SUMIF 公式，实为 {v!r}"
        )


def test_matrix_ratio_uses_iserror_guard(wb):
    """比例行应有 ISERROR 兜底（分母 0 → 0）。"""
    ws = wb["函证结果汇总表E0-1"]
    for r in (31, 33, 34):  # 三行比例
        v = ws.cell(r, 5).value  # 取 E 列样本
        assert v and "ISERROR" in str(v), f"R{r} E 列缺 ISERROR 兜底: {v!r}"


# ── 7. E0-1 下区四块锚点与固定文字 ───────────────────────────────────────────

# 四块锚点（反向自检：数量锚，防工具函数读到空集合恒绿）
E0_1_LOWER_ZONE_BLOCKS = {
    "C27": "一、函证情况",
    "O27": "二、样本选择",
    "V27": "三、审计说明",
    "V34": "四、审计结论",
}


def test_lower_zone_block_anchors(wb):
    """下区四块锚点逐字。"""
    ws = wb["函证结果汇总表E0-1"]
    for ref, expected in E0_1_LOWER_ZONE_BLOCKS.items():
        actual = ws[ref].value
        assert actual == expected, f"{ref}: {actual!r} != {expected!r}"
    # 反向自检：确实有 4 块
    assert len(E0_1_LOWER_ZONE_BLOCKS) == 4


def test_lower_zone_sample_selection_texts(wb):
    """「二、样本选择」三段固定文字逐字（O28 / O29 / O30+O31 合并）。"""
    ws = wb["函证结果汇总表E0-1"]
    assert ws["O28"].value == "所有银行账户全部函证（包括零余额账户和在本期内注销的账户）。"
    assert ws["O29"].value == "如果存在未函证的银行账户应记录不执行函证程序的理由。"
    # O30+O31 是一句话被源模板拆成两格
    o30 = str(ws["O30"].value or "")
    o31 = str(ws["O31"].value or "")
    combined = o30 + o31
    assert "审计准则规定的可以不执行银行函证程序的理由是" in combined
    assert "对财务报表不重要且与之相关的重大错报风险很低" in combined


def test_lower_zone_o30_o31_split_is_not_complete_sentence(wb):
    """反向自检：O30 单独不是完整句子（证明合并渲染的必要性）。"""
    ws = wb["函证结果汇总表E0-1"]
    o30 = str(ws["O30"].value or "")
    assert not o30.rstrip().endswith("。"), (
        "O30 不应以句号结尾（它是半句话，与 O31 合并才完整）"
    )


def test_lower_zone_audit_note_three_items(wb):
    """「三、审计说明」三条小标题锚点（V28/V30/V32+V33）。"""
    ws = wb["函证结果汇总表E0-1"]
    assert ws["V28"].value == "1.对询证函保持的控制的说明"
    assert ws["V30"].value == "2.针对不符事项的程序"
    # V32+V33 同样是一句话被拆成两格
    v32 = str(ws["V32"].value or "")
    v33 = str(ws["V33"].value or "")
    combined_v = v32 + v33
    assert "3.如果银行回函中存在未函证的其他信息" in combined_v
    assert "未函证信息的影响，并考虑实施进一步审计程序" in combined_v


def test_lower_zone_v32_typo_preserved(wb):
    """V32 的源模板笔误「本函证证」原样保留（防被"顺手修正"后三向比对打红）。"""
    ws = wb["函证结果汇总表E0-1"]
    v32 = str(ws["V32"].value or "")
    assert "本函证证" in v32, "源模板 V32 应含笔误「本函证证」"


def test_lower_zone_v32_v33_split_is_not_complete_sentence(wb):
    """反向自检：V32 单独不是完整句子。"""
    ws = wb["函证结果汇总表E0-1"]
    v32 = str(ws["V32"].value or "")
    # V32 以"应考虑"结尾（后接 V33），不是完整句子
    assert not v32.rstrip().endswith("。") and not v32.rstrip().endswith("程序"), (
        "V32 不应是完整句子"
    )


def test_lower_zone_tips(wb):
    """「提示」区 A37 标题 + A38 四条（合并区 A38:N38）。"""
    ws = wb["函证结果汇总表E0-1"]
    assert ws["A37"].value == "提示：对收到的回函重点检查："
    a38 = str(ws["A38"].value or "")
    assert "1.确认与发出的询证函是否为同一份、是否为原件" in a38
    assert "2.核实回函快递物流信息" in a38
    assert "3.关注回函中是否包含免责或其他限制性条款" in a38
    assert "4.检查函证信息是否相符" in a38


def test_lower_zone_a38_is_merged(wb):
    """A38 合并区 A38:N38（防未来拆分后只读到 A38 一格）。"""
    ws = wb["函证结果汇总表E0-1"]
    merged_refs = [str(mc) for mc in ws.merged_cells.ranges if mc.min_row == 38]
    assert any("A38" in r and "N38" in r for r in merged_refs), (
        f"A38:N38 合并区缺失。当前 R38 合并区: {merged_refs}"
    )


# ── 8. E0-7 跟函函证过程控制（五段话术 + 三个「是否」+ 签名栏）──────────────


def test_e0_7_tab_name_is_exact(wb):
    """tab 名逐字：`跟函函证过程控制E0-7`（两个「函」字）。"""
    assert "跟函函证过程控制E0-7" in wb.sheetnames


def test_e0_7_tab_name_typo_guard(wb):
    """反向自检：少写一个「函」必不命中。"""
    assert "跟函证过程控制E0-7" not in wb.sheetnames
    assert "跟函函证过程E0-7" not in wb.sheetnames


def test_e0_7_five_scenario_keywords(wb):
    """五段话术关键词（对公柜台/XX部门/寄回致同/公示内容一致/补记）。"""
    ws = wb["跟函函证过程控制E0-7"]
    # 把全部非空格子拼成一个大字符串做关键词搜索
    all_text = "\n".join(
        str(ws.cell(r, c).value or "")
        for r in range(1, ws.max_row + 1)
        for c in range(1, ws.max_column + 1)
        if ws.cell(r, c).value
    )
    # 五段场景的判据关键词
    assert "对公柜台办理" in all_text, "缺 场景1：对公柜台办理"
    assert "至银行" in all_text and "部门" in all_text, "缺 场景2：至银行XX部门"
    assert "寄回致同会计师事务所" in all_text, "缺 场景3：留函待寄回"
    assert "公示内容一致" in all_text, "缺 场景4：已对全部项目作出回应+公示核对"
    assert "收回" in all_text, "缺 场景5：事后收回补记"


def test_e0_7_has_employee_number(wb):
    """话术含「工号」占位（银行工作人员+复核人员各一个）。"""
    ws = wb["跟函函证过程控制E0-7"]
    all_text = "\n".join(
        str(ws.cell(r, c).value or "")
        for r in range(1, ws.max_row + 1)
        for c in range(1, ws.max_column + 1)
        if ws.cell(r, c).value
    )
    assert "工号" in all_text, "E0-7 话术须含「工号」占位"


def test_e0_7_three_confirmation_questions(wb):
    """三个「是否」项逐字。"""
    ws = wb["跟函函证过程控制E0-7"]
    questions = []
    for r in range(1, ws.max_row + 1):
        v = ws.cell(r, 1).value
        if v and "是否" in str(v):
            questions.append(str(v).strip())
    expected = [
        "是否了解处理函证的通常流程和处理人员",
        "是否确认询证函处理人员的身份及权限",
        "处理人员是否按正常流程处理",
    ]
    assert questions == expected, f"三个「是否」项不匹配: {questions}"


def test_e0_7_has_signature_line(wb):
    """签名栏存在。"""
    ws = wb["跟函函证过程控制E0-7"]
    found = False
    for r in range(1, ws.max_row + 1):
        v = ws.cell(r, 1).value
        if v and "签名" in str(v):
            found = True
            break
    assert found, "E0-7 缺签名栏"

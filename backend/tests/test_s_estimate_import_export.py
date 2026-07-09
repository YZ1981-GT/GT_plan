"""S 类计算型底稿导入导出三级端点 — 单元测试

验证：
- 导出模板/导出数据/导入数据三端点注册
- service 层 _col_letter / get_sheet_filename 辅助函数
- _SHEET_CONFIGS 配置完整性
- RFC5987 中文文件名编码
- 多区块分 sheet 导出逻辑

Requirements: 10.1, 10.2, 10.3, 10.4
Spec: .kiro/specs/s-estimate-calculation-workpapers/ Task 6.3
"""

import io
from unittest.mock import AsyncMock, patch
from urllib.parse import quote

import pytest


class TestEndpointRegistration:
    """三端点注册验证"""

    def test_export_template_endpoint_exists(self):
        """export-template 端点已注册."""
        from app.routers.s_estimate_calculation import router

        routes = [r.path for r in router.routes]
        assert any("export-template" in r for r in routes)

    def test_export_data_endpoint_exists(self):
        """export-data 端点已注册."""
        from app.routers.s_estimate_calculation import router

        routes = [r.path for r in router.routes]
        assert any("export-data" in r for r in routes)

    def test_import_data_endpoint_exists(self):
        """import-data 端点已注册."""
        from app.routers.s_estimate_calculation import router

        routes = [r.path for r in router.routes]
        assert any("import-data" in r for r in routes)

    def test_export_template_is_get(self):
        """export-template 为 GET 方法."""
        from app.routers.s_estimate_calculation import router

        for route in router.routes:
            if "export-template" in route.path:
                assert "GET" in route.methods
                break
        else:
            pytest.fail("export-template route not found")

    def test_export_data_is_get(self):
        """export-data 为 GET 方法."""
        from app.routers.s_estimate_calculation import router

        for route in router.routes:
            if "export-data" in route.path:
                assert "GET" in route.methods
                break
        else:
            pytest.fail("export-data route not found")

    def test_import_data_is_post(self):
        """import-data 为 POST 方法."""
        from app.routers.s_estimate_calculation import router

        for route in router.routes:
            if "import-data" in route.path:
                assert "POST" in route.methods
                break
        else:
            pytest.fail("import-data route not found")


class TestSheetConfigs:
    """Sheet 配置完整性"""

    def test_three_sheets_configured(self):
        """三个 sheet 已配置."""
        from app.services.s_estimate_import_export_service import _SHEET_CONFIGS

        assert "S21-2" in _SHEET_CONFIGS
        assert "S20-unrelated" in _SHEET_CONFIGS
        assert "S20-noSubstance" in _SHEET_CONFIGS

    def test_s21_2_is_grid_type(self):
        """S21-2 为 grid 类型."""
        from app.services.s_estimate_import_export_service import _SHEET_CONFIGS

        assert _SHEET_CONFIGS["S21-2"]["type"] == "grid"

    def test_s20_sheets_are_dynamic_rows(self):
        """S20 sheets 为 dynamic_rows 类型."""
        from app.services.s_estimate_import_export_service import _SHEET_CONFIGS

        assert _SHEET_CONFIGS["S20-unrelated"]["type"] == "dynamic_rows"
        assert _SHEET_CONFIGS["S20-noSubstance"]["type"] == "dynamic_rows"

    def test_s21_2_has_14_headers(self):
        """S21-2 表头: 成本类目 + 12月 + 合计 = 14列."""
        from app.services.s_estimate_import_export_service import _SHEET_CONFIGS

        headers = _SHEET_CONFIGS["S21-2"]["headers"]
        assert len(headers) == 14
        assert headers[0] == "成本类目"
        assert headers[1] == "1月"
        assert headers[12] == "12月"
        assert headers[13] == "合计"

    def test_s20_has_4_headers(self):
        """S20 明细表头: 4列."""
        from app.services.s_estimate_import_export_service import _SHEET_CONFIGS

        assert len(_SHEET_CONFIGS["S20-unrelated"]["headers"]) == 4
        assert len(_SHEET_CONFIGS["S20-noSubstance"]["headers"]) == 4

    def test_s21_categories_count(self):
        """S21-2 有7个成本类目."""
        from app.services.s_estimate_import_export_service import _S21_CATEGORIES

        assert len(_S21_CATEGORIES) == 7
        assert "采购成本" in _S21_CATEGORIES
        assert "安全管理费" in _S21_CATEGORIES


class TestHelperFunctions:
    """辅助函数测试"""

    def test_col_letter_single(self):
        """列号转字母 — 单字母."""
        from app.services.s_estimate_import_export_service import _col_letter

        assert _col_letter(1) == "A"
        assert _col_letter(26) == "Z"

    def test_col_letter_double(self):
        """列号转字母 — 双字母."""
        from app.services.s_estimate_import_export_service import _col_letter

        assert _col_letter(27) == "AA"
        assert _col_letter(28) == "AB"

    def test_get_sheet_filename_s21(self):
        """S21-2 文件名生成."""
        from app.services.s_estimate_import_export_service import get_sheet_filename

        result = get_sheet_filename("S21-2", "模板")
        assert result == "S21-2 开发支出资本化分析_模板.xlsx"
        assert result.endswith(".xlsx")

    def test_get_sheet_filename_s20_unrelated(self):
        """S20-unrelated 文件名生成."""
        from app.services.s_estimate_import_export_service import get_sheet_filename

        result = get_sheet_filename("S20-unrelated", "数据")
        assert result == "S20 与主营无关收入明细_数据.xlsx"

    def test_get_sheet_filename_s20_nosubstance(self):
        """S20-noSubstance 文件名生成."""
        from app.services.s_estimate_import_export_service import get_sheet_filename

        result = get_sheet_filename("S20-noSubstance", "模板")
        assert result == "S20 不具备商业实质收入明细_模板.xlsx"

    def test_get_sheet_filename_fallback(self):
        """未知 sheet 使用 fallback."""
        from app.services.s_estimate_import_export_service import get_sheet_filename

        result = get_sheet_filename("UNKNOWN", "模板")
        assert "UNKNOWN" in result


class TestRFC5987Encoding:
    """RFC5987 中文文件名编码验证（Req 10.4）"""

    def test_chinese_filename_url_encoded(self):
        """中文文件名通过 urllib.parse.quote 编码."""
        filename = "S21-2 开发支出资本化分析_模板.xlsx"
        encoded = quote(filename)
        # 编码后不应包含中文字符
        assert "开发" not in encoded
        # 但应可以解码回来
        from urllib.parse import unquote
        assert unquote(encoded) == filename

    def test_content_disposition_format(self):
        """Content-Disposition header 格式正确（RFC5987）."""
        filename = "S20 与主营无关收入明细_数据.xlsx"
        encoded = quote(filename)
        header = f"attachment; filename*=UTF-8''{encoded}"
        assert header.startswith("attachment; filename*=UTF-8''")
        assert "S20" in header


class TestExportTemplate:
    """导出模板功能测试"""

    @pytest.mark.asyncio
    async def test_export_template_returns_bytes(self):
        """导出模板返回 BytesIO."""
        from app.services.s_estimate_import_export_service import export_template

        db = AsyncMock()
        buffer = await export_template("test-wp-id", db, sheet="S21-2")
        assert isinstance(buffer, io.BytesIO)
        # 验证是有效的 xlsx（zip header: PK）
        content = buffer.getvalue()
        assert content[:2] == b"PK"

    @pytest.mark.asyncio
    async def test_export_template_all_sheets(self):
        """不传 sheet 时导出全部 sheet."""
        from app.services.s_estimate_import_export_service import export_template

        db = AsyncMock()
        buffer = await export_template("test-wp-id", db, sheet=None)
        assert isinstance(buffer, io.BytesIO)

        from openpyxl import load_workbook as lw
        wb = lw(buffer, read_only=True)
        # 应有 3 个 sheet（S21-2 / S20-unrelated / S20-noSubstance）
        assert len(wb.sheetnames) == 3
        wb.close()

    @pytest.mark.asyncio
    async def test_export_template_single_sheet(self):
        """传 sheet 时仅导出指定 sheet."""
        from app.services.s_estimate_import_export_service import export_template

        db = AsyncMock()
        buffer = await export_template("test-wp-id", db, sheet="S21-2")

        from openpyxl import load_workbook as lw
        wb = lw(buffer, read_only=True)
        assert len(wb.sheetnames) == 1
        assert "S21-2" in wb.sheetnames[0]
        wb.close()

    @pytest.mark.asyncio
    async def test_s21_template_has_category_rows(self):
        """S21-2 模板预填 7 个类目名称行头."""
        from app.services.s_estimate_import_export_service import export_template, _S21_CATEGORIES

        db = AsyncMock()
        buffer = await export_template("test-wp-id", db, sheet="S21-2")

        from openpyxl import load_workbook as lw
        wb = lw(buffer)
        ws = wb.active
        # 第 2~8 行 A 列应为类目名
        for i, cat in enumerate(_S21_CATEGORIES, 2):
            assert ws.cell(row=i, column=1).value == cat
        wb.close()


class TestMultiBlockExport:
    """多区块分 sheet 导出验证（Req 10.3）"""

    @pytest.mark.asyncio
    async def test_all_sheets_present_in_full_export(self):
        """全量导出包含所有已配置的 sheet."""
        from app.services.s_estimate_import_export_service import export_template, _SHEET_CONFIGS

        db = AsyncMock()
        buffer = await export_template("test-wp-id", db, sheet=None)

        from openpyxl import load_workbook as lw
        wb = lw(buffer, read_only=True)
        assert len(wb.sheetnames) == len(_SHEET_CONFIGS)
        wb.close()

# -*- coding: utf-8 -*-
"""Task 53 守卫 —— K 循环（管理循环）Excel 独立 entry 迁移。

spec: workpaper-html-onlyoffice-bidirectional-writeback-closure · Wave 5 Task 53
Requirements: 6.1, 6.2, 6.6, 12.1, 12.4, 12.8, 12.9, 12.10, 12.11, 12.12, 14.1, 1.4
Properties: 20（部分）· 24（部分）· 69 · 70

═══ 判据设计（继承 Task 51/52 已确立的做法，逐条说明本轮的差异）═══

**① source_ref 三边锁 + 第四边**。声明 → openpyxl 真读源 xlsx → impl 常量现读。
K 循环额外要求「消费边 = statement-position import 路径字面量」，且**排除落在双引号
字符串内的匹配** —— `workpaperSyncLegacyBaseline.generated.ts` 的 JSON ``"snippet"``
字段里含完整的 ``from '...'`` 形态，会骗过路径字面量口径。禁符号名 grep：本轮实测
`useK9DualMode.ts` 只在**文档注释**里提到 `useWorkpaperEntryDualMode`，符号名口径会把它
算成消费方，于是「K 对共享基类贡献 0 条边」这个结论会翻。

**② orphan 判定用可达性而非入度**。K 循环没有二阶 orphan，但这不能靠「没找到」下结论 ——
判据是 `components/workpaper/composables/index.ts` **不存在**（现算），故该目录下的模块
没有 barrel 可躲。两侧都验：7 个一阶 orphan 边数各 0 **且** barrel 文件确实不存在。

**③ 剥注释必须保留行号**（同长空白替换）。K 循环这条尤其关键：11 个宿主的第二处
``el-segmented`` 落在**块注释里**，剥注释后消失。若用删行式剥注释，`#Lnn` 坐标会整体上移
⇒ 全部 toolbar 锚点假红。

**④ 计数一律从来源节现算等值**。分母为空的 Property **明确不宣称通过**
（见 slice 的 `property_denominators`）。

**⑤ 存量欠账用 `pytest.mark.xfail(strict=True)`，禁 `pytest.skip`**。skip 记为通过 = fail-open。

═══ 本轮**不得照抄 J 循环判据**的四处（照抄必假红）═══

* J: `J0` 两个解析函数都返回 None（J 无函证册）→ **K: `K0` 都返回真实文件**（K 有函证册）。
* J: 共享基类消费方 29 → 26（J 贡献 3 条）→ **K: 29 → 29（K 贡献 0 条）**。
* J: `hardcoded_scan_result` 六个模式全 0 → **K: 第六个模式命中 13 处**（披露表行 id 模板）。
* J: 有 `parent_duplicate_summary` 与 2 个 orphan barrel → **K: 两者都没有**（0 条 / 无 barrel）。

═══ 反向自检 ═══

`TestGuardSelfChecks` 里每个 helper 都有「故意写错必失败」的双向自检：剥注释真的隐藏了
注释内容且**行号不变**、双引号字符串检测器两侧都对、路径解析是路径式而非 stem 式、
toolbar 区块截取有界、位置化 token 两种形态都抓、家族判别式对「同时含熵与兜底」的那处
归 family_b（不是放过）、openpyxl 读 sheet 名不 strip。
"""
from __future__ import annotations

import hashlib
import importlib.util
import json
import os
import pathlib
import re
import sys
from typing import Any

import pytest
from openpyxl import load_workbook

_THIS = pathlib.Path(__file__).resolve()
ROOT = _THIS.parents[3]
BACKEND = ROOT / "backend"
FRONTEND = ROOT / "audit-platform" / "frontend" / "src"
WP_COMPONENTS = FRONTEND / "components" / "workpaper"
WP_COMPOSABLES = WP_COMPONENTS / "composables"
CYCLE_COMPOSABLES = FRONTEND / "composables" / "workpaper"
SYNC_DIR = WP_COMPONENTS / "sync"
DATA = BACKEND / "data"

MANIFEST_SLICE_PATH = DATA / "workpaper_sync_k_cycle_manifest_slice.json"
DELETION_PLAN_PATH = DATA / "workpaper_sync_k_cycle_deletion_plan.json"
PARADIGM_PATH = DATA / "workpaper_sync_migration_paradigm.json"
FULL_MANIFEST_PATH = DATA / "workpaper_sync_entry_manifest.json"
OVERLAY_PATH = DATA / "workpaper_sync_entry_overlay.json"
SIBLING_SLICE_PATHS = {
    "D": DATA / "workpaper_sync_d_cycle_manifest_slice.json",
    "E": DATA / "workpaper_sync_e_cycle_manifest_slice.json",
    "F": DATA / "workpaper_sync_f_cycle_manifest_slice.json",
    "G": DATA / "workpaper_sync_g_cycle_manifest_slice.json",
    "H": DATA / "workpaper_sync_h_cycle_manifest_slice.json",
    "I": DATA / "workpaper_sync_i_cycle_manifest_slice.json",
    "J": DATA / "workpaper_sync_j_cycle_manifest_slice.json",
}
CONTRACT_DIR = DATA / "workpaper_sync_contracts"
TEMPLATE_DIR = BACKEND / "wp_templates"
TEMPLATE_INDEX = TEMPLATE_DIR / "_index.json"
K_TEMPLATE_DIR = TEMPLATE_DIR / "K"
CHECKLIST_ROUTER = BACKEND / "app" / "routers" / "checklist_responses.py"
REGISTRY = BACKEND / "app" / "services" / "workpaper_sync" / "adapters" / "registry.py"
HTML_RENDERER_REGISTRY = WP_COMPONENTS / "htmlRendererRegistry.ts"
SHARED_BASE = WP_COMPOSABLES / "useWorkpaperEntryDualMode.ts"
LEGACY_BASELINE_GENERATED = SYNC_DIR / "workpaperSyncLegacyBaseline.generated.ts"
CHECKLIST_PERSISTENCE = CYCLE_COMPOSABLES / "useChecklistPersistence.ts"
COMPOSABLES_BARREL = WP_COMPOSABLES / "index.ts"

NOTICE_MODULE = SYNC_DIR / "workpaperEntrySyncNotice.ts"
NOTICE_COMPONENT = SYNC_DIR / "GtEntrySyncCapabilityNotice.vue"
NOTICE_COMPONENT_NAME = "GtEntrySyncCapabilityNotice"

K_CYCLE_SPECS_PY = BACKEND / "app" / "services" / "four_table" / "k_cycle_specs.py"
REPORT_LINE_ACCOUNTS_PY = BACKEND / "app" / "services" / "four_table" / "report_line_accounts.py"
RENDER_STRATEGY_DIR = BACKEND / "app" / "routers" / "wp_render_strategies"

WRITEOFF_OWNER = WP_COMPOSABLES / "useK1WriteoffCheck.ts"
WRITEOFF_TAB = WP_COMPONENTS / "k1" / "inspection" / "K1TabWriteoffCheck.vue"

CAPABILITY_ENUM = ("bidirectional", "single_html", "single_onlyoffice", "unreachable")
HTML_COUNTERPART_VERDICTS = ("none", "exists")

#: 四条 pilot 契约的**实证**归属（逐文件读顶层 `review.entry_id`）。
#: 🔴 文件名与 entry_id 不同构，禁按文件名猜。
PILOT_CONTRACT_OWNERS = {
    "xlsx/b60/gt-b60-bundle",
    "xlsx/gt-d2-accounts-receivable",
    "xlsx/gt-g7-long-term-equity-main",
    "xlsx/gt-h1-fixed-assets",
}

#: 共享基类的窄口径（statement-position）消费方总数。K 循环贡献 0 条 ⇒ 前后都是这个数。
SHARED_BASE_CONSUMERS = 29
#: 宽口径（逐文件 token 命中）文件数。差集证明窄口径不是空谈。
SHARED_BASE_WIDE_FILES = 33


# ════════════════════════════════════════════════════════════════════════════
# helper（每个都有反向自检，见 TestGuardSelfChecks）
# ════════════════════════════════════════════════════════════════════════════
def _load(path: pathlib.Path) -> dict:
    return json.loads(path.read_text(encoding="utf-8"))


def _sha256_of(path: pathlib.Path) -> str:
    return hashlib.sha256(path.read_bytes()).hexdigest()


def _blank_keep_newlines(match: "re.Match[str]") -> str:
    """把匹配段替换成同长空白，**保留换行** —— 行号绝不能变。"""
    return re.sub(r"[^\n]", " ", match.group(0))


def _strip_comments(source: str) -> str:
    """剥块注释 / 行注释 / HTML 注释，行号保持不变。

    🔴 K 循环这条尤其关键：11 个宿主的第二处 `el-segmented` 落在块注释里，
    剥注释后消失；若用删行式剥注释，全部 `#Lnn` 锚点会整体上移 ⇒ 假红。
    """
    source = re.sub(r"/\*.*?\*/", _blank_keep_newlines, source, flags=re.S)
    source = re.sub(r"<!--.*?-->", _blank_keep_newlines, source, flags=re.S)
    source = re.sub(
        r"(?<![:\w\"'`\\])//[^\n]*", lambda m: " " * len(m.group(0)), source
    )
    return source


def _vue_template(source: str) -> str:
    """取 `<template>` 顶层块（不被内层 `<template v-if>` 提前截断），行号保留。"""
    start = source.find("<template>")
    if start < 0:
        return ""
    end = source.rfind("</template>")
    if end < 0:
        return ""
    head = re.sub(r"[^\n]", " ", source[:start])
    return head + source[start:end + len("</template>")]


def _toolbar_block(template: str, gate_anchor_line: int, span: int = 12) -> str:
    """从 `ui_toolbar_gate.anchor` 那一行起截 span 行 —— **有界**，不许扫全文件。"""
    lines = template.split("\n")
    lo = max(0, gate_anchor_line - 1)
    return "\n".join(lines[lo:lo + span])


def _cells_of(path: pathlib.Path, sheet: str, cells: str) -> list[Any]:
    """openpyxl 真读某区间。🔴 sheet 名**不得 strip**。"""
    wb = load_workbook(path, read_only=True, data_only=True)
    try:
        if sheet not in wb.sheetnames:
            raise AssertionError(
                f"{path.name} 里没有 sheet {sheet!r}（现有：{wb.sheetnames}）"
            )
        ws = wb[sheet]
        out: list[Any] = []
        for row in ws[cells]:
            for c in row:
                out.append(c.value)
        return out
    finally:
        wb.close()


def _sheet_names(path: pathlib.Path) -> list[str]:
    wb = load_workbook(path, read_only=True, data_only=True)
    try:
        return list(wb.sheetnames)
    finally:
        wb.close()


def _indexed_relpaths() -> set[str]:
    idx = _load(TEMPLATE_INDEX)
    return {str(f["relative_path"]).replace("\\", "/") for f in idx["files"]}


def _resolve_repo(rel: str) -> pathlib.Path:
    return ROOT / rel.split("#L")[0]


def _line_no_of(ref: str) -> int:
    m = re.search(r"#L(\d+)$", ref)
    assert m, f"{ref!r} 不含 #L<行号>"
    return int(m.group(1))


def _line_at(ref: str) -> str:
    path = _resolve_repo(ref)
    lines = path.read_text(encoding="utf-8").split("\n")
    n = _line_no_of(ref)
    assert 1 <= n <= len(lines), f"{ref} 行号越界（文件共 {len(lines)} 行）"
    return lines[n - 1]


_IMPORT_FORMS = (
    re.compile(r"""from\s*['"]([^'"\n]+)['"]"""),
    re.compile(r"""import\s*\(\s*['"]([^'"\n]+)['"]"""),
    re.compile(r"""vi\.mock\s*\(\s*['"]([^'"\n]+)['"]"""),
    re.compile(r"""require\s*\(\s*['"]([^'"\n]+)['"]"""),
)


def _inside_double_quoted_string(line: str, pos: int) -> bool:
    """`pos` 是否落在**双引号字符串**内。

    🔴 `workpaperSyncLegacyBaseline.generated.ts` 的 JSON `"snippet"` 字段里含完整的
    ``from '...'`` 形态；不排除它，「K 对共享基类贡献 0 条边」这个结论会翻。
    """
    depth = 0
    i = 0
    while i < pos and i < len(line):
        ch = line[i]
        if ch == "\\":
            i += 2
            continue
        if ch == '"':
            depth ^= 1
        i += 1
    return bool(depth)


_FE_CACHE: dict[pathlib.Path, str] = {}


def _cached_text(path: pathlib.Path) -> str:
    if path not in _FE_CACHE:
        _FE_CACHE[path] = path.read_text(encoding="utf-8", errors="replace")
    return _FE_CACHE[path]


def _all_frontend() -> list[pathlib.Path]:
    return [
        p for p in FRONTEND.rglob("*")
        if p.is_file() and p.suffix in (".ts", ".vue", ".tsx", ".js")
    ]


def _resolve_spec(spec: str, importer: pathlib.Path) -> pathlib.Path | None:
    """把 import spec 解析成仓库内路径（**路径式**，不是 stem 式）。"""
    if spec.startswith("@/"):
        return FRONTEND / spec[2:]
    if spec.startswith("."):
        return (importer.parent / spec).resolve()
    return None


def _statement_edges_to(target: pathlib.Path) -> tuple[list[str], list[str]]:
    """返回 (生产边, 测试边)，只认 statement-position 且排除双引号字符串内的匹配。"""
    prod: list[str] = []
    test: list[str] = []
    stem = target.with_suffix("")
    for f in _all_frontend():
        text = _cached_text(f)
        if target.stem not in text:
            continue
        for i, line in enumerate(text.split("\n"), 1):
            for rx in _IMPORT_FORMS:
                for m in rx.finditer(line):
                    if _inside_double_quoted_string(line, m.start()):
                        continue
                    r = _resolve_spec(m.group(1), f)
                    if r is None:
                        continue
                    if r in (stem, target) or r.with_suffix("") == stem:
                        ref = f"{f.relative_to(ROOT).as_posix()}#L{i}"
                        (test if _is_test_path(f) else prod).append(ref)
    return prod, test


def _is_test_path(path: pathlib.Path) -> bool:
    p = path.as_posix()
    return "__tests__" in p or p.endswith((".spec.ts", ".test.ts"))


def _wide_scope_edge_files(token: str) -> set[str]:
    """宽口径：逐文件 token 命中即算（用于证明窄口径不是空谈）。"""
    out: set[str] = set()
    for f in _all_frontend():
        if token in _cached_text(f):
            out.add(f.relative_to(ROOT).as_posix())
    return out


def _k_cycle_files() -> list[pathlib.Path]:
    """K 域生产文件（口径与 slice 的 `scanned_scope` 逐字一致）。"""
    out: list[pathlib.Path] = []
    for n in range(1, 14):
        d = WP_COMPONENTS / f"k{n}"
        if d.is_dir():
            for p in d.rglob("*"):
                if p.is_file() and p.suffix in (".ts", ".vue") and "__tests__" not in p.as_posix():
                    out.append(p)
    for p in sorted(WP_COMPONENTS.iterdir()):
        if p.suffix == ".vue" and re.match(r"^GtK\d+", p.name):
            out.append(p)
    for p in sorted(WP_COMPOSABLES.iterdir()):
        if p.suffix == ".ts" and re.match(r"^(use)?[kK](1[0-3]|[1-9])(?![0-9])", p.name):
            out.append(p)
    for n in range(1, 14):
        d = CYCLE_COMPOSABLES / f"k{n}"
        if d.is_dir():
            for p in d.rglob("*.ts"):
                if "__tests__" not in p.as_posix():
                    out.append(p)
    return sorted(set(out))


_IDENTITY_KEY = re.compile(r"(?<![\w$])(rowId|rowKey|id)\s*:\s*([^,\n]+)")
_POSITIONAL_TOKEN = re.compile(r"(?:^|[^\w$])(?:i|idx|index)(?:\s*\+\s*1)?(?:\s*\}|\s*[,)\]`]|$)")
_POSITIONAL_INTERP = re.compile(r"\$\{\s*(?:i|idx|index)\s*\}")
_ENTROPY = re.compile(r"Date\.now\(\)|Math\.random\(\)")
_FALLBACK = re.compile(r"\?\?|\|\|")
_DISPLAY_SEQ_SITE = re.compile(r"\bseq\s*:\s*[^,]*(?:idx|index|\bi\b)\s*\+\s*1")

_HARDCODED_PATTERNS = {
    "blankRows_literal_count": re.compile(r"blankRows\s*\([^,)]*,\s*\d+\s*\)"),
    "horizontal_company_column_literals": re.compile(r"['\"]公司\s*[1-9]\d*['\"]"),
    "column_key_uses_label": re.compile(r"key\s*:\s*(?:row|col|c)\.label"),
    "row_cell_key_uses_label": re.compile(r"cells\s*\[\s*(?:row|r)\.label\s*\]"),
    "seed_placeholder_literal": re.compile(r"['\"](?:项目|成本项目|单位)\s*N['\"]"),
    "positional_row_id_template": re.compile(
        r"`(?:seed|row|detail|item)-\$\{\s*(?:i|idx|index)\s*\}"
    ),
}


def _positional_identity_hits(files: list[pathlib.Path]) -> list[tuple[str, str, str]]:
    """返回 [(ref, key, value)]，与 slice 的 scan_recipe 同一口径。"""
    hits: list[tuple[str, str, str]] = []
    for p in files:
        src = _strip_comments(_cached_text(p))
        for i, line in enumerate(src.split("\n"), 1):
            for m in _IDENTITY_KEY.finditer(line):
                val = m.group(2).strip()
                if _POSITIONAL_TOKEN.search(val) or _POSITIONAL_INTERP.search(val):
                    hits.append((f"{p.relative_to(ROOT).as_posix()}#L{i}", m.group(1), val))
    return hits


def _family_of(value: str) -> str:
    """三族判别式 —— 可复算的布尔表达式，不是人工归类。

    🔴 顺序要紧：`K7TabDisclosureSoe.vue#L401` 同时含熵与兜底 ⇒ 必须归 family_b（是缺陷）。
    """
    if _ENTROPY.search(value) and not _FALLBACK.search(value):
        return "c"
    if _FALLBACK.search(value):
        return "b"
    return "a"


def _hardcoded_hits(files: list[pathlib.Path], name: str) -> list[str]:
    rx = _HARDCODED_PATTERNS[name]
    out: list[str] = []
    for p in files:
        src = _strip_comments(_cached_text(p))
        for i, line in enumerate(src.split("\n"), 1):
            if rx.search(line):
                out.append(f"{p.relative_to(ROOT).as_posix()}#L{i}")
    return out


def _display_seq_sites(files: list[pathlib.Path]) -> list[str]:
    out: list[str] = []
    for p in files:
        src = _strip_comments(_cached_text(p))
        for i, line in enumerate(src.split("\n"), 1):
            if _DISPLAY_SEQ_SITE.search(line):
                out.append(f"{p.relative_to(ROOT).as_posix()}#L{i}")
    return out


def _const_literal(path: pathlib.Path, name: str) -> tuple[str, int]:
    """现读 `const NAME = '<literal>'`，返回 (字面量, 行号)。剥注释后再找。"""
    src = _strip_comments(path.read_text(encoding="utf-8"))
    m = re.search(r"^const %s\s*=\s*'([^']+)'" % re.escape(name), src, re.M)
    assert m, f"{path.name} 里找不到 `const {name} = '...'`"
    return m.group(1), src[:m.start()].count("\n") + 1


def _exact_literal_hits(files: list[pathlib.Path], literal: str) -> list[str]:
    """精确字面量命中（整个引号内容 == literal），每行至多计一次。"""
    out: list[str] = []
    for p in files:
        src = _strip_comments(_cached_text(p))
        for i, line in enumerate(src.split("\n"), 1):
            for m in re.finditer(r"""['"]([^'"\n]+)['"]""", line):
                if m.group(1) == literal:
                    out.append(f"{p.relative_to(ROOT).as_posix()}#L{i}")
                    break
    return out


def _wp_code_of(entry: dict) -> str:
    return "K" + re.search(r"K(\d+)", entry["wp_code_pattern"]).group(1)


# ════════════════════════════════════════════════════════════════════════════
# fixtures
# ════════════════════════════════════════════════════════════════════════════
@pytest.fixture(scope="module")
def manifest_slice() -> dict:
    return _load(MANIFEST_SLICE_PATH)


@pytest.fixture(scope="module")
def deletion_plan() -> dict:
    return _load(DELETION_PLAN_PATH)


@pytest.fixture(scope="module")
def paradigm() -> dict:
    return _load(PARADIGM_PATH)


@pytest.fixture(scope="module")
def full_manifest() -> dict:
    return _load(FULL_MANIFEST_PATH)


@pytest.fixture(scope="module")
def k_files() -> list[pathlib.Path]:
    return _k_cycle_files()


@pytest.fixture(scope="module")
def k1_template() -> pathlib.Path:
    p = K_TEMPLATE_DIR / "K1 其他应收款.xlsx"
    assert p.exists(), f"权威模板不存在：{p}"
    return p


# ════════════════════════════════════════════════════════════════════════════
class TestGuardSelfChecks:
    """helper 的反向自检：故意写错必失败。没有这一节，下面所有判据都可能是空跑。"""

    def test_strip_comments_hides_content_but_keeps_line_numbers(self) -> None:
        src = "a\n/* el-segmented\nstill comment */\nb // el-segmented\nc\n"
        out = _strip_comments(src)
        assert "el-segmented" not in out, "剥注释没生效"
        assert len(out.split("\n")) == len(src.split("\n")), (
            "🔴 剥注释改变了行号 —— 全部 #Lnn 锚点会整体上移变成假红"
        )
        assert out.split("\n")[0] == "a" and out.split("\n")[4] == "c"

    def test_strip_comments_really_hides_the_k_host_commented_segmented(self) -> None:
        """K 循环的实况判据：11 个宿主剥注释前有 2 处 el-segmented、剥后 1 处。"""
        host = WP_COMPONENTS / "GtK11AssetImpairmentLoss.vue"
        raw = host.read_text(encoding="utf-8")
        before = [i for i, l in enumerate(raw.split("\n"), 1) if "el-segmented" in l]
        after = [
            i for i, l in enumerate(_strip_comments(raw).split("\n"), 1)
            if "el-segmented" in l
        ]
        assert len(before) == 2, f"前提变了：剥注释前应有 2 处，实得 {before}"
        assert after == [before[0]], (
            f"剥注释后应只剩第一处（toolbar 那处），实得 {after}"
        )

    def test_inside_double_quoted_string_detector_works_both_ways(self) -> None:
        line = '  "snippet": "import x from \'./useK1DualMode\'",'
        pos = line.index("from")
        assert _inside_double_quoted_string(line, pos) is True, (
            "🔴 检测器没识别出 JSON snippet 内的伪边"
        )
        real = "import x from './useK1DualMode'"
        assert _inside_double_quoted_string(real, real.index("from")) is False

    def test_resolve_spec_is_path_based_not_stem_based(self) -> None:
        importer = WP_COMPONENTS / "GtK1OtherReceivables.vue"
        got = _resolve_spec("./composables/useK1DualMode", importer)
        assert got == WP_COMPOSABLES / "useK1DualMode"
        # stem 式会把这个也算成命中；路径式必须不同
        other = _resolve_spec("./k1/useK1DualMode", importer)
        assert other != got, "🔴 解析退化成 stem 相等 —— 会把同名不同目录的模块误判成消费方"

    def test_vue_template_extraction_does_not_stop_at_inner_template(self) -> None:
        src = "<template>\n<div>\n<template v-if=\"x\">\n<p/>\n</template>\n</div>\n</template>\n"
        out = _vue_template(src)
        assert out.count("</template>") == 2, "内层 template 把顶层块提前截断了"
        assert len(out.split("\n")) <= len(src.split("\n"))

    def test_toolbar_block_is_bounded(self) -> None:
        template = "\n".join(f"L{i}" for i in range(1, 60))
        block = _toolbar_block(template, 10, span=5)
        assert block.split("\n") == ["L10", "L11", "L12", "L13", "L14"]

    def test_positional_token_catches_both_interp_and_plus_one(self) -> None:
        assert _POSITIONAL_INTERP.search("`row-${idx}`")
        assert _POSITIONAL_TOKEN.search("rows[i + 1]") or _POSITIONAL_TOKEN.search("x, i + 1)")
        assert not _POSITIONAL_INTERP.search("`row-${uuid}`")

    def test_family_discriminator_puts_entropy_plus_fallback_in_b(self) -> None:
        """🔴 反向自检：同时含熵与兜底的那处必须归 family_b（是缺陷），不是放过。"""
        assert _family_of("String(raw?.id || `grant-${idx}-${Date.now()}`)") == "b"
        assert _family_of("`row-${idx}-${Date.now()}`") == "c"
        assert _family_of("`row-${idx}`") == "a"
        assert _family_of("raw.rowKey ?? `row-${idx}`") == "b"

    def test_cells_of_does_not_strip_the_sheet_name(self, k1_template: pathlib.Path) -> None:
        names = _sheet_names(k1_template)
        target = next(n for n in names if "K1-9" in n)
        assert _cells_of(k1_template, target, "A1:A1") is not None
        with pytest.raises(AssertionError):
            _cells_of(k1_template, target + " ", "A1:A1")

    def test_const_literal_reads_from_source_not_from_comment(self) -> None:
        value, line = _const_literal(WRITEOFF_OWNER, "ITEM_ID")
        assert value == "K1-9-writeoff"
        assert _line_at(f"{WRITEOFF_OWNER.relative_to(ROOT).as_posix()}#L{line}").strip() == (
            "const ITEM_ID = 'K1-9-writeoff'"
        )
        with pytest.raises(AssertionError):
            _const_literal(WRITEOFF_OWNER, "NO_SUCH_CONSTANT")

    def test_k_cycle_denominator_is_non_vacuous(self, k_files: list[pathlib.Path]) -> None:
        assert len(k_files) >= 300, (
            f"K 域文件分母塌了（{len(k_files)} 个）—— 下面所有穷举判据都会变成空跑"
        )
        assert any("GtK1OtherReceivables.vue" in p.as_posix() for p in k_files)
        assert any("useK1WriteoffCheck.ts" in p.as_posix() for p in k_files)
        assert any("/k11/" in p.as_posix() for p in k_files), (
            "🔴 前缀正则漏了 K10..K13（负向断言 (?![0-9]) 写错时会这样）"
        )

    def test_all_required_artifacts_exist(self) -> None:
        for p in (
            MANIFEST_SLICE_PATH, DELETION_PLAN_PATH, PARADIGM_PATH, FULL_MANIFEST_PATH,
            OVERLAY_PATH, CONTRACT_DIR, TEMPLATE_INDEX, K_TEMPLATE_DIR, CHECKLIST_ROUTER,
            REGISTRY, HTML_RENDERER_REGISTRY, SHARED_BASE, CHECKLIST_PERSISTENCE,
            NOTICE_MODULE, NOTICE_COMPONENT, K_CYCLE_SPECS_PY, REPORT_LINE_ACCOUNTS_PY,
            RENDER_STRATEGY_DIR, WRITEOFF_OWNER, WRITEOFF_TAB,
        ):
            assert p.exists(), f"判据依赖的产物不存在：{p}"
        for name, p in SIBLING_SLICE_PATHS.items():
            assert p.exists(), f"兄弟 slice 不存在：{name} -> {p}"


# ════════════════════════════════════════════════════════════════════════════
class TestSliceScopeIsRecomputable:
    @staticmethod
    def _k_prefixed(full_manifest: dict) -> list[dict]:
        out = []
        for e in full_manifest["entries"]:
            pats = (e.get("wp_match") or {}).get("wp_code_patterns") or []
            if any(str(p).startswith("K") for p in pats):
                out.append(e)
        return out

    def test_selection_rule_recomputes_the_entry_set(
        self, manifest_slice: dict, full_manifest: dict
    ) -> None:
        recomputed = {
            e["entry_id"] for e in self._k_prefixed(full_manifest)
            if e["document_type"] == "xlsx" and e["independent_entry"]
        }
        declared = {e["entry_id"] for e in manifest_slice["independent_entries"]}
        assert declared == recomputed, (
            "slice 的 entry 集合与按 selection_rule 现算的集合不等 —— "
            f"多写 {sorted(declared - recomputed)}，漏写 {sorted(recomputed - declared)}"
        )
        assert len(recomputed) == 13

    def test_scope_counters_match_the_manifest(
        self, manifest_slice: dict, full_manifest: dict
    ) -> None:
        scope = manifest_slice["slice_scope"]
        k_all = self._k_prefixed(full_manifest)
        assert scope["k_prefixed_entry_total"] == len(k_all)
        assert scope["k_prefixed_independent_total"] == sum(
            1 for e in k_all if e["independent_entry"])
        assert scope["independent_entry_count"] == len(manifest_slice["independent_entries"])
        assert scope["cycle"] == "K" and scope["document_type"] == "xlsx"
        assert {e["document_type"] for e in k_all} == {"xlsx"}, (
            "excluded_from_slice 声称「K 循环的非 xlsx entry」排除集为空，现算却不是"
        )

    def test_parent_duplicate_count_is_really_zero_both_ways(
        self, manifest_slice: dict, full_manifest: dict
    ) -> None:
        """🔴 KD-5：0 条必须两侧都验，且本 slice 顶层不得有 parent_duplicate_summary。"""
        k_ids = {e["entry_id"] for e in manifest_slice["independent_entries"]}
        children = [
            e["entry_id"] for e in full_manifest["entries"]
            if e.get("parent_entry_id") in k_ids
        ]
        assert children == [], f"manifest 里出现了 K 的 parent_duplicate 子入口：{children}"
        assert manifest_slice["slice_scope"]["parent_duplicate_count"] == 0
        overlay = _load(OVERLAY_PATH)
        k_rules = [
            r for r in overlay["parent_rules"]
            if re.search(r"/k\d+/", str(r.get("file_glob", "")))
        ]
        assert k_rules == [], f"overlay 的 parent_rules 里有 K 的 file_glob：{k_rules}"
        assert "parent_duplicate_summary" not in manifest_slice, (
            "🔴 触发条件不成立却写了 parent_duplicate_summary —— 那就是 additive 死声明"
        )

    def test_currency_variant_section_is_absent_because_untriggered(
        self, manifest_slice: dict
    ) -> None:
        assert "currency_variant_model" not in manifest_slice
        assert "why_no_currency_variant_model_section" in (
            manifest_slice["paradigm_schema_conflict"]["new_section_not_in_schema"]
        )

    def test_host_module_edges_in_the_renderer_registry_are_real(
        self, manifest_slice: dict
    ) -> None:
        """AC 12.9 的「入口可达」用**模块边**判，不按符号名 grep。"""
        reg = _strip_comments(HTML_RENDERER_REGISTRY.read_text(encoding="utf-8"))
        for entry in manifest_slice["independent_entries"]:
            stem = pathlib.Path(entry["host_path"]).stem
            needle = f"import('./{stem}.vue')"
            assert needle in reg, (
                f"{entry['entry_id']} 的宿主在 htmlRendererRegistry 里没有模块边 "
                f"（找 {needle!r}）⇒ 不能声称「入口可达」"
            )

    def test_no_pilot_contract_belongs_to_the_k_cycle(self, manifest_slice: dict) -> None:
        """契约归属**逐文件读 review.entry_id**，不数文件个数、不按文件名猜。

        🔴 口径：只对**生产契约**（`review_status != "candidate"`）要求 entry_id 非空。
        `_example.candidate.json` 是 step 6「review_status=candidate 不得注册生产 adapter」
        的反例分母，它的 review.entry_id 本就是 null。
        """
        iso = manifest_slice["cross_entry_isolation"]
        k_ids = {e["entry_id"] for e in manifest_slice["independent_entries"]}
        owners: set[str] = set()
        prod: list[str] = []
        cand: list[str] = []
        files = sorted(CONTRACT_DIR.glob("*.json"))
        assert files, "契约目录为空 —— 分母塌了，本判据变成空跑"
        for p in files:
            doc = _load(p)
            if doc.get("review_status") == "candidate":
                cand.append(p.name)
                assert (doc.get("review") or {}).get("entry_id") is None, (
                    f"{p.name} 是 candidate 却带 entry_id ⇒ 可能被误注册成生产契约"
                )
                continue
            prod.append(p.name)
            assert doc.get("review_status") == "reviewed", (
                f"{p.name}: 生产契约的 review_status 必须是 reviewed（step 6）"
            )
            owner = (doc.get("review") or {}).get("entry_id")
            assert owner, f"{p.name} 的 review.entry_id 缺失"
            owners.add(owner)
            assert owner not in k_ids, f"{p.name} 的 review.entry_id={owner} 属本 slice"
        assert owners == PILOT_CONTRACT_OWNERS, (
            f"四条 pilot 契约的实证归属变了：期望 {sorted(PILOT_CONTRACT_OWNERS)}，"
            f"实得 {sorted(owners)}"
        )
        assert prod == iso["production_contract_files"], "生产契约清单登记与磁盘不符"
        assert cand == iso["candidate_contract_files"], "candidate 契约清单登记与磁盘不符"
        assert cand, "🔴 candidate 反例分母为空 ⇒ 「candidate 不得进生产」这条判据没有对象"

    def test_k0_confirmation_workbook_exists_but_has_no_entry(
        self, manifest_slice: dict, full_manifest: dict
    ) -> None:
        """🔴 KD-4：K0 有册但无 entry，归 Task 57。与 J0（无册）相反。"""
        hits = [e["entry_id"] for e in full_manifest["entries"] if "k0" in e["entry_id"].lower()]
        assert hits == [], f"manifest 里出现了 K0 entry：{hits}"
        k0 = K_TEMPLATE_DIR / "K0 管理循环函证.xlsx"
        assert k0.exists(), "K0 册不存在 —— 与 KD-4 的登记矛盾"
        rec = next(
            f for f in manifest_slice["authoritative_templates"]["files"]
            if f["name"] == k0.name
        )
        assert rec["belongs_to_entry"] is None and rec["excluded_reason"]
        excl = manifest_slice["slice_scope"]["excluded_from_slice"]
        assert any("K0" in str(x["what"]) for x in excl)

    def test_gtkam_host_is_declared_out_of_scope(
        self, manifest_slice: dict, full_manifest: dict
    ) -> None:
        """「按文件名前缀 GtK 枚举宿主」这条捷径被显式否决。"""
        gtk = sorted(
            p.name for p in WP_COMPONENTS.iterdir()
            if p.suffix == ".vue" and p.name.startswith("GtK")
        )
        hosts = {pathlib.Path(e["host_path"]).name for e in manifest_slice["independent_entries"]}
        extra = set(gtk) - hosts
        assert extra == {"GtKamWorkpaper.vue"}, (
            f"GtK* 文件集与宿主集的差集变了：{sorted(extra)}"
        )
        for e in full_manifest["entries"]:
            if pathlib.Path(e["host_path"]).name == "GtKamWorkpaper.vue":
                pats = (e.get("wp_match") or {}).get("wp_code_patterns") or []
                assert not any(re.match(r"^K\d", str(p)) for p in pats), (
                    f"GtKamWorkpaper 的 wp_code_patterns 命中了 K+数字：{pats}"
                )
        assert any(
            "GtKamWorkpaper" in str(x["what"])
            for x in manifest_slice["slice_scope"]["excluded_from_slice"]
        )


# ════════════════════════════════════════════════════════════════════════════
class TestAdjudicationLegality:
    def test_every_entry_has_a_binary_html_counterpart_verdict(self, manifest_slice: dict) -> None:
        for e in manifest_slice["independent_entries"]:
            assert e["html_counterpart_verdict"] in HTML_COUNTERPART_VERDICTS, (
                f"{e['entry_id']}: verdict={e['html_counterpart_verdict']!r} 不是二值之一"
                "（AP-3：unresolved/unknown/空值都不是结论）"
            )
            assert e["html_counterpart_source_refs"], f"{e['entry_id']}: source_refs 为空"

    def test_single_onlyoffice_requires_no_html_counterpart(self, manifest_slice: dict) -> None:
        for e in manifest_slice["independent_entries"]:
            if e["capability"] == "single_onlyoffice":
                assert e["html_counterpart_verdict"] == "none", (
                    f"{e['entry_id']}: 裁 single_onlyoffice 却有 HTML 对端（AC 12.8）"
                )

    def test_adjudication_reason_is_not_circular(
        self, manifest_slice: dict, paradigm: dict
    ) -> None:
        """AP-1：不得用「本任务应交付的产物尚不存在」当裁 single 的理由。"""
        ap1 = next(
            a for a in paradigm["adjudication_criteria"]["anti_patterns"] if a["id"] == "AP-1"
        )
        markers = ap1["circular_reason_markers"]
        for e in manifest_slice["independent_entries"]:
            if e["capability"] != "single_onlyoffice":
                continue
            blob = json.dumps(e["adjudication"], ensure_ascii=False)
            hit = [m for m in markers if m in blob]
            assert not hit, f"{e['entry_id']}: 裁决理由是循环论证（命中 {hit}）"

    def test_capability_is_null_and_pending_fields_are_complete(
        self, manifest_slice: dict, paradigm: dict
    ) -> None:
        """SR-3 右支：13/13 待裁决，三字段按语义齐备。"""
        schema = paradigm["slice_schema"]
        fields = schema["required_pending_verdict_fields"]
        semantics = schema["pending_verdict_field_semantics"]
        bp_ids = {b["id"] for b in manifest_slice["blocking_preconditions"]}
        pending = 0
        for e in manifest_slice["independent_entries"]:
            if e["capability"] is not None:
                assert e["capability"] in CAPABILITY_ENUM
                continue
            pending += 1
            for f in fields:
                kind = semantics[f]
                v = e.get(f)
                if kind == "non_empty_string":
                    assert isinstance(v, str) and v.strip(), f"{e['entry_id']}.{f}"
                elif kind == "member_of_capability_enum":
                    assert v in CAPABILITY_ENUM, f"{e['entry_id']}.{f}={v!r}"
                elif kind == "non_empty_list":
                    assert isinstance(v, list) and v, f"{e['entry_id']}.{f}"
                else:
                    pytest.fail(f"语义 kind 未知：{kind}")
            unknown = set(e["capability_target_blocked_by"]) - bp_ids
            assert not unknown, (
                f"{e['entry_id']}: capability_target_blocked_by 引用了不存在的阻断项 {sorted(unknown)}"
            )
        assert pending == 13
        assert manifest_slice["honest_adjudication_summary"]["slice_counters"][
            "unadjudicated"
        ] == pending, "SR-9：待裁决条数与计数器不等"

    def test_capability_matches_honest_capability(self, manifest_slice: dict) -> None:
        for e in manifest_slice["independent_entries"]:
            assert e["capability"] == e["adjudication"]["honest_capability"], (
                f"{e['entry_id']}: SR-4 双口径"
            )

    def test_pending_entries_carry_no_identity(
        self, manifest_slice: dict, paradigm: dict
    ) -> None:
        fields = paradigm["slice_schema"]["presence_required_value_may_be_null"]
        for e in manifest_slice["independent_entries"]:
            for f in fields:
                assert f in e, f"{e['entry_id']}: `{f}` 键不得缺（值可为 null）"
                assert e[f] is None, (
                    f"{e['entry_id']}: 待裁决态却挂着 `{f}`={e[f]!r}"
                )

    def test_adjudication_carries_both_negative_reasons(self, manifest_slice: dict) -> None:
        for e in manifest_slice["independent_entries"]:
            adj = e["adjudication"]
            assert adj["not_single_html_because"].strip()
            assert adj["not_bidirectional_because"].strip()
            assert adj["reason"].strip()

    def test_every_blocking_precondition_is_referenced_by_someone(
        self, manifest_slice: dict
    ) -> None:
        """反向：登记了阻断项却没有任何 entry 指向它 = 死声明。"""
        declared = {b["id"] for b in manifest_slice["blocking_preconditions"]}
        used: set[str] = set()
        for e in manifest_slice["independent_entries"]:
            used |= set(e["capability_target_blocked_by"])
        orphaned = declared - used
        assert not orphaned, f"这些阻断项没有任何 entry 引用（死声明）：{sorted(orphaned)}"
        assert used <= declared

    def test_all_blocking_preconditions_carry_task48_fields(
        self, manifest_slice: dict, paradigm: dict
    ) -> None:
        extra = paradigm["slice_schema"]["effective_from_task_48"]["required_blocking_precondition"]
        for b in manifest_slice["blocking_preconditions"]:
            for f in extra:
                assert f in b and b[f], f"{b['id']}: 缺 Task 48 起必填的 `{f}`"
            assert b.get("consequence") or b.get("observable_consequences"), (
                f"{b['id']}: 必须写明后果"
            )

    def test_bp4_is_the_writeoff_blocker_with_a_must_fix_before(
        self, manifest_slice: dict
    ) -> None:
        """tasks.md 正文第二条的落点：writeoff 必须 adapter 化 ⇒ BP-4 存在且只挂在 K1 上。"""
        bp4 = next(b for b in manifest_slice["blocking_preconditions"] if b["id"] == "BP-4")
        assert "adapter" in bp4["must_fix_before"] or "step 9" in bp4["must_fix_before"]
        holders = [
            _wp_code_of(e) for e in manifest_slice["independent_entries"]
            if "BP-4" in e["capability_target_blocked_by"]
        ]
        assert holders == ["K1"], (
            f"BP-4（writeoff）应只挂在 K1 上（writeoff 表是 K1-9），实得 {holders}"
        )

    def test_manifest_mirror_divergence_is_registered_not_silently_equal(
        self, manifest_slice: dict, full_manifest: dict
    ) -> None:
        """🔴 overlay 陷阱：slice 与 manifest_mirror **必须不一致**且已登记为 BP-9。"""
        by_id = {e["entry_id"]: e for e in full_manifest["entries"]}
        overlay = _load(OVERLAY_PATH)
        default = overlay["defaults_by_component"]["GtOnlyOfficeSheet"]
        diverged = 0
        for e in manifest_slice["independent_entries"]:
            mirror = e["manifest_mirror"]
            live = by_id[e["entry_id"]]
            assert mirror["capability"] == live["capability"], (
                f"{e['entry_id']}: manifest_mirror.capability 与 manifest 现值不符 —— 镜像抄错了"
            )
            assert mirror["html_store"] == live["html_store"]
            assert mirror["capability"] == default["capability"], (
                f"{e['entry_id']}: 镜像值应等于 overlay 默认值（证明它不是逐 entry 裁决）"
            )
            assert mirror["html_store"] == default["html_store"]
            if mirror["capability"] != e["capability"]:
                diverged += 1
                assert mirror["why_not_adopted"], f"{e['entry_id']}: 分歧未说明"
                assert "BP-9" in mirror["why_not_adopted"]
        assert diverged == 13, (
            "🔴 slice 与 manifest 镜像必须**全部**不一致（overlay 默认 single_onlyoffice vs "
            f"slice 待裁决），实得 {diverged}/13 —— 断言相等就是把 overlay 默认值当裁决"
        )
        assert any(b["id"] == "BP-9" for b in manifest_slice["blocking_preconditions"])

    def test_ac15_is_declared_not_applicable_and_ac14_is_the_live_one(
        self, manifest_slice: dict
    ) -> None:
        s = manifest_slice["honest_adjudication_summary"]
        assert s["ac_15_not_applicable_because"].strip()
        assert "AC 1.4" in s["ac_15_not_applicable_because"], (
            "必须写明「AC 1.5 不适用、生效的是 AC 1.4」"
        )
        for e in manifest_slice["independent_entries"]:
            cap = e["capability"]
            assert not (isinstance(cap, str) and cap.startswith("single_")), (
                "AC 1.5 的触发条件（被裁为纯 OO / 纯 HTML）不成立才能声称不适用"
            )
            assert e["adapter_id"] is None
            assert e["migration_state"] == "legacy_fake_bidirectional", (
                "AC 1.4 的触发条件（未注册 adapter）必须成立"
            )


# ════════════════════════════════════════════════════════════════════════════
class TestHtmlCounterpartIsSourceBacked:
    def test_source_refs_point_at_real_paths_and_lines(self, manifest_slice: dict) -> None:
        for e in manifest_slice["independent_entries"]:
            for ref in e["html_counterpart_source_refs"]:
                p = _resolve_repo(ref)
                assert p.exists(), f"{e['entry_id']}: source_ref 指向不存在的文件 {ref}"
                if "#L" in ref:
                    _line_at(ref)  # 越界会 assert

    def test_html_store_endpoint_exists_in_the_router(self, manifest_slice: dict) -> None:
        router = _strip_comments(CHECKLIST_ROUTER.read_text(encoding="utf-8"))
        assert re.search(r'@router\.put\(""', router), (
            "checklist_responses 路由缺 PUT ⇒ 「HTML 对端存在」这个结论没有服务端落点"
        )
        assert re.search(r'@router\.get\(""', router)
        for e in manifest_slice["independent_entries"]:
            assert e["html_counterpart"]["store"] == "checklist_responses"
            assert e["html_counterpart"]["row_identity_key"] == "item_id"

    def test_every_host_imports_the_shared_persistence_adapter(
        self, manifest_slice: dict
    ) -> None:
        """13 条 entry 的持久化通道同构 —— 这是「不设 transport_key_resolution 节」的判据。"""
        prod, _ = _statement_edges_to(CHECKLIST_PERSISTENCE)
        hosts = {e["host_path"] for e in manifest_slice["independent_entries"]}
        covered = {ref.split("#L")[0] for ref in prod}
        missing = hosts - covered
        assert not missing, (
            f"这些宿主没有 import useChecklistPersistence ⇒ 同构性主张不成立：{sorted(missing)}"
        )
        assert manifest_slice["paradigm_schema_conflict"]["new_section_not_in_schema"][
            "why_no_transport_key_resolution_section"
        ]
        assert "transport_key_resolution" not in manifest_slice

    def test_template_ref_resolves_through_the_runtime_index(self, manifest_slice: dict) -> None:
        indexed = _indexed_relpaths()
        for e in manifest_slice["independent_entries"]:
            ref = e["template_ref"]
            p = K_TEMPLATE_DIR / ref["workbook"]
            assert p.exists(), f"{e['entry_id']}: 权威模板不存在 {p}"
            assert _sha256_of(p) == ref["sha256"], f"{e['entry_id']}: 模板 digest 漂移"
            assert len(_sheet_names(p)) == ref["sheet_count"]
            assert f"K/{ref['workbook']}" in indexed, (
                f"{e['entry_id']}: 模板不在 _index.json 的运行时索引里"
            )
            assert ref["in_runtime_index"] is True

    def test_each_entry_has_its_own_workbook(self, manifest_slice: dict) -> None:
        books = [e["template_ref"]["workbook"] for e in manifest_slice["independent_entries"]]
        assert len(set(books)) == len(books) == 13, (
            "两条 entry 共用一本权威模板 ⇒ Property 70 的模板维度失守"
        )

    def test_entry_profile_fields_mirror_the_manifest(
        self, manifest_slice: dict, full_manifest: dict
    ) -> None:
        by_id = {e["entry_id"]: e for e in full_manifest["entries"]}
        for e in manifest_slice["independent_entries"]:
            live = by_id[e["entry_id"]]
            assert e["editability"] == live["editability"]
            assert e["room_model"] == live["room_model"]
            assert e["canonical_resolver"] == live["canonical_resolver"]
            assert e["mount_count"] == len(live["mounts"])
            assert e["host_path"] == live["host_path"]
            assert e["wp_code_pattern"] in (live.get("wp_match") or {}).get(
                "wp_code_patterns", [])


# ════════════════════════════════════════════════════════════════════════════
class TestOrphanDualModeInventory:
    @staticmethod
    def _modules(manifest_slice: dict) -> list[dict]:
        return manifest_slice["orphan_dual_mode_inventory"]["modules"]

    def test_declared_orphans_really_have_no_reachability(self, manifest_slice: dict) -> None:
        for m in self._modules(manifest_slice):
            p = ROOT / m["file"]
            assert p.exists(), f"{m['id']}: 声明的 orphan 文件不存在 {m['file']}"
            prod, test = _statement_edges_to(p)
            assert prod == [], f"{m['id']}: 声称 orphan 却有生产边 {prod}"
            assert test == [], f"{m['id']}: 声称 orphan 却有测试边 {test}"
            assert m["production_consumers"] == 0 and m["test_only_consumers"] == 0
            assert len(p.read_text(encoding="utf-8").split("\n")) == m["lines"], (
                f"{m['id']}: 行数登记与磁盘不符"
            )

    def test_orphans_are_all_first_order_because_there_is_no_barrel(
        self, manifest_slice: dict
    ) -> None:
        """🔴 二阶恒 0 的判据是「barrel 不存在」，不是「没找到」。"""
        assert not COMPOSABLES_BARREL.exists(), (
            f"{COMPOSABLES_BARREL} 出现了 —— 二阶 orphan 从此可能存在，判据必须重写"
        )
        inv = manifest_slice["orphan_dual_mode_inventory"]
        assert inv["summary"]["orphan_dual_mode_second_order"] == 0
        assert inv["summary"]["orphan_barrels"] == 0
        for m in self._modules(manifest_slice):
            assert m["orphan_order"] == "first_order"
            assert m["barrel_only_reachable_via"] is None

    def test_live_modules_have_exactly_one_edge_to_the_declared_host(
        self, manifest_slice: dict
    ) -> None:
        live = manifest_slice["orphan_dual_mode_inventory"]["live_modules"]
        assert len(live) == 6
        hosts = {e["host_path"] for e in manifest_slice["independent_entries"]}
        for m in live:
            p = ROOT / m["file"]
            prod, test = _statement_edges_to(p)
            assert len(prod) == 1, f"{m['id']}: 在用 composable 应恰 1 条生产边，实得 {prod}"
            assert prod == m["production_consumers"], f"{m['id']}: 边登记与现算不符"
            assert test == [], f"{m['id']}: 出现了测试边 {test}"
            assert prod[0].split("#L")[0] in hosts

    def test_the_thirteen_dual_mode_files_split_exactly_seven_plus_six(
        self, manifest_slice: dict
    ) -> None:
        files = sorted(
            p for p in WP_COMPOSABLES.iterdir()
            if re.fullmatch(r"useK(1[0-3]|[1-9])DualMode\.ts", p.name)
        )
        assert len(files) == 13, f"useK*DualMode.ts 文件数变了：{[p.name for p in files]}"
        inv = manifest_slice["orphan_dual_mode_inventory"]
        declared = {ROOT / m["file"] for m in inv["modules"]}
        declared |= {ROOT / m["file"] for m in inv["live_modules"]}
        assert declared == set(files), (
            "orphan + live 的并集必须恰覆盖 13 个文件（交集为空）"
        )
        assert len(inv["modules"]) == 7 and len(inv["live_modules"]) == 6

    def test_orphan_and_host_twin_prefix_relationship_recomputes(
        self, manifest_slice: dict
    ) -> None:
        """K4/K5/K6 撞前缀 3 处、K1/K2/K3/K7 孪生无持久化 4 处；3 + 4 == 7。"""
        collide = 0
        no_persist = 0
        for m in self._modules(manifest_slice):
            code = m["wp_code"]
            host = ROOT / next(
                e["host_path"] for e in manifest_slice["independent_entries"]
                if _wp_code_of(e) == code
            )
            src = _strip_comments(host.read_text(encoding="utf-8"))
            found = re.findall(r"DUAL_MODE_STORAGE_PREFIX\s*=\s*'([^']+)'", src)
            twin = found[0] if found else None
            assert m["twin_localStorage_prefix"] == twin, (
                f"{m['id']}: 孪生前缀登记 {m['twin_localStorage_prefix']!r} != 现读 {twin!r}"
            )
            own, _ = _const_literal(ROOT / m["file"], "STORAGE_PREFIX")
            assert m["localStorage_prefix"] == own
            if twin == own:
                collide += 1
                assert m["prefix_collision_with_twin"] is True
            else:
                assert m["prefix_collision_with_twin"] is False
            if twin is None:
                no_persist += 1
        s = manifest_slice["orphan_dual_mode_inventory"]["summary"]
        assert collide == s["orphans_with_prefix_collision_with_host_twin"] == 3
        assert no_persist == s["orphans_whose_twin_has_no_persistence"] == 4
        assert collide + no_persist == len(self._modules(manifest_slice)) == 7

    def test_legacy_endpoint_direct_calls_are_registered_per_class(
        self, manifest_slice: dict
    ) -> None:
        """step 6 的违反面：health 全 13 个 composable 都调，config 只有 6 个在用的调。

        🔴 端点扫描必须**认反引号模板字面量** —— config 端点写成
        `` `/api/workpapers/${wpId.value}/sheets/${encodeURIComponent(sn)}/onlyoffice-config` ``，
        只认单/双引号的正则会把 6 处 config 直调漏成 0（首轮实测）。
        """
        inv = manifest_slice["orphan_dual_mode_inventory"]
        assert "反引号" in inv["endpoint_scan_recipe"], (
            "endpoint_scan_recipe 必须写明反引号模板字面量口径"
        )
        health = 0
        config = 0
        for m in inv["modules"] + inv["live_modules"]:
            src = _strip_comments((ROOT / m["file"]).read_text(encoding="utf-8"))
            has_health = "/api/workpapers/onlyoffice/health" in src
            has_config = "onlyoffice-config" in src
            assert has_health == any(
                "onlyoffice/health" in e for e in m["legacy_endpoints_called"]), m["id"]
            assert has_config == any(
                "onlyoffice-config" in e for e in m["legacy_endpoints_called"]), m["id"]
            health += int(has_health)
            config += int(has_config)
        s = inv["summary"]
        assert health == s["modules_calling_legacy_health_endpoint"] == 13
        assert config == s["modules_calling_legacy_config_endpoint"] == 6, (
            "🔴 config 计数必须与 health 分开：orphan 侧只调 health"
        )

    def test_hosts_calling_the_legacy_health_endpoint_directly(
        self, manifest_slice: dict
    ) -> None:
        n = 0
        for e in manifest_slice["independent_entries"]:
            src = _strip_comments((ROOT / e["host_path"]).read_text(encoding="utf-8"))
            direct = "/api/workpapers/onlyoffice/health" in src
            inline = e["dual_mode_carrier"]["kind"] == "host_inlined_iife"
            assert direct == inline, (
                f"{e['entry_id']}: 「宿主直调 health」与「载体是内联 IIFE」应等价，"
                f"实得 direct={direct} inline={inline}"
            )
            n += int(direct)
        assert n == manifest_slice["orphan_dual_mode_inventory"]["summary"][
            "hosts_calling_legacy_health_endpoint_directly"
        ] == 7

    def test_shared_base_counts_are_recomputed_both_ways(self, manifest_slice: dict) -> None:
        """🔴 KD-3：K 贡献 0 条 ⇒ 29 前后不变。窄口径 vs 宽口径两侧都验。"""
        prod, test = _statement_edges_to(SHARED_BASE)
        narrow = prod + test
        assert len(narrow) == SHARED_BASE_CONSUMERS, (
            f"共享基类窄口径消费方数变了：{len(narrow)}（期望 {SHARED_BASE_CONSUMERS}）"
        )
        k_edges = [r for r in narrow if re.search(r"/(GtK\d|k\d+/|useK\d)", r)]
        assert k_edges == [], f"🔴 K 循环竟然有共享基类的边：{k_edges}"
        wide = _wide_scope_edge_files("useWorkpaperEntryDualMode")
        assert len(wide) == SHARED_BASE_WIDE_FILES, (
            f"宽口径文件数变了：{len(wide)}（期望 {SHARED_BASE_WIDE_FILES}）"
        )
        narrow_files = {r.split("#L")[0] for r in narrow}
        diff = wide - narrow_files
        assert "audit-platform/frontend/src/components/workpaper/composables/useK9DualMode.ts" in diff, (
            "🔴 useK9DualMode.ts 只在注释里提到基类名，必须落在宽窄差集里 —— "
            "否则「符号名口径会误判」这个论证是空的"
        )
        assert LEGACY_BASELINE_GENERATED.relative_to(ROOT).as_posix() in diff
        sb = manifest_slice["orphan_dual_mode_inventory"]["shared_base"]
        assert sb["statement_position_consumers"] == SHARED_BASE_CONSUMERS
        assert sb["k_cycle_contribution"] == 0
        assert sb["remaining_after_k_cycle_work"] == SHARED_BASE_CONSUMERS
        assert sb["preserved"] is True

    def test_orphan_summary_counts_recompute(self, manifest_slice: dict) -> None:
        inv = manifest_slice["orphan_dual_mode_inventory"]
        s = inv["summary"]
        assert s["orphan_dual_mode_total"] == len(inv["modules"])
        assert s["orphan_dual_mode_first_order"] == sum(
            1 for m in inv["modules"] if m["orphan_order"] == "first_order")
        assert s["live_dual_mode_total"] == len(inv["live_modules"])
        assert s["orphan_lines_total"] == sum(m["lines"] for m in inv["modules"])
        assert s["counting_note"].strip()


# ════════════════════════════════════════════════════════════════════════════
class TestKCycleFormDifferences:
    @staticmethod
    def _diff(manifest_slice: dict, did: str) -> dict:
        return next(
            d for d in manifest_slice["k_cycle_form_differences"]["differences"]
            if d["id"] == did
        )

    def test_all_seven_differences_carry_a_recompute_recipe(self, manifest_slice: dict) -> None:
        diffs = manifest_slice["k_cycle_form_differences"]["differences"]
        assert len(diffs) == 7
        assert [d["id"] for d in diffs] == [f"KD-{i}" for i in range(1, 8)]
        for d in diffs:
            assert d["what"].strip() and d["why_it_matters"].strip() and d["recompute"].strip()

    def test_kd2_carrier_split_is_real_and_exclusive(self, manifest_slice: dict) -> None:
        measured = self._diff(manifest_slice, "KD-2")["measured"]
        inline_n = 0
        dedicated_n = 0
        for code, rec in measured.items():
            host = ROOT / next(
                e["host_path"] for e in manifest_slice["independent_entries"]
                if _wp_code_of(e) == code
            )
            src = _strip_comments(host.read_text(encoding="utf-8"))
            iife = [
                i for i, l in enumerate(src.split("\n"), 1)
                if re.search(r"const dualMode = \(\(\) =>", l)
            ]
            imp = [
                i for i, l in enumerate(src.split("\n"), 1)
                if re.search(r"import \{ use%sDualMode \}" % code, l)
            ]
            assert not (iife and imp), f"{code}: 两种载体同时出现 —— 集合不再互斥"
            if iife:
                inline_n += 1
                assert rec["inline_iife_site"] == (
                    f"{host.relative_to(ROOT).as_posix()}#L{iife[0]}"
                )
                assert rec["imports_dedicated_composable"] is None
            else:
                dedicated_n += 1
                assert rec["inline_iife_site"] is None
                assert rec["imports_dedicated_composable"] == (
                    f"{host.relative_to(ROOT).as_posix()}#L{imp[0]}"
                )
        assert inline_n == 7 and dedicated_n == 6

    def test_kd6_two_gate_shapes_coexist(self, manifest_slice: dict) -> None:
        """🔴 判据不得统一要求 v-if —— K8/K9 走 disabled 选项范式，会假红。"""
        measured = self._diff(manifest_slice, "KD-6")["measured"]
        with_vif = set(measured["hosts_with_v_if_gate"])
        with_disabled = set(measured["hosts_with_disabled_option_gate"])
        assert with_vif & with_disabled == set()
        assert len(with_vif) + len(with_disabled) == 13
        assert with_disabled == {"K8", "K9"}, f"disabled 门控的宿主集变了：{with_disabled}"
        for e in manifest_slice["independent_entries"]:
            code = _wp_code_of(e)
            src = _strip_comments((ROOT / e["host_path"]).read_text(encoding="utf-8"))
            sites = [
                i for i, l in enumerate(src.split("\n"), 1)
                if 'v-if="dualMode.isOoAvailable.value"' in l
            ]
            assert sites == e["ui_toolbar_gate"]["second_level_gate_sites"]
            if code in with_disabled:
                assert sites == []
                comp = WP_COMPOSABLES / f"use{code}DualMode.ts"
                csrc = _strip_comments(comp.read_text(encoding="utf-8"))
                assert "disabled: !isOoAvailable.value" in csrc, (
                    f"{code}: 声称走 disabled 范式，composable 里却找不到该门控"
                )
            else:
                assert len(sites) == 1

    def test_kd7_commented_segmented_would_fool_a_raw_grep(self, manifest_slice: dict) -> None:
        measured = self._diff(manifest_slice, "KD-7")["measured"]
        fooled = 0
        for e in manifest_slice["independent_entries"]:
            code = _wp_code_of(e)
            raw = (ROOT / e["host_path"]).read_text(encoding="utf-8")
            before = [i for i, l in enumerate(raw.split("\n"), 1) if "el-segmented" in l]
            after = [
                i for i, l in enumerate(_strip_comments(raw).split("\n"), 1)
                if "el-segmented" in l
            ]
            assert before == measured[code]["before"], f"{code}: 剥前行号变了"
            assert after == measured[code]["after"], f"{code}: 剥后行号变了"
            assert len(after) == 1, f"{code}: 剥注释后应恰 1 处 toolbar 分段器"
            assert set(after) <= set(before)
            if len(before) > len(after):
                fooled += 1
        assert fooled >= 10, (
            f"🔴 只有 {fooled} 个宿主会被裸 grep 骗到 —— 前提变了，KD-7 的论证需要重算"
        )

    def test_kd4_template_count_exceeds_entry_count_by_the_k0_book(
        self, manifest_slice: dict
    ) -> None:
        files = manifest_slice["authoritative_templates"]["files"]
        assert len(files) == 14
        owned = [f for f in files if f["belongs_to_entry"]]
        unowned = [f for f in files if not f["belongs_to_entry"]]
        assert len(owned) == 13 and len(unowned) == 1
        assert unowned[0]["name"].startswith("K0 ")


# ════════════════════════════════════════════════════════════════════════════
class TestWriteoffCapability:
    """tasks.md 正文第二条：「纯客户端 writeoff 能力必须 adapter 化或单模式裁决」。

    本节的结论是 **writeoff 不是纯客户端** ⇒ 只能取「必须 adapter 化」这一支（BP-4）。
    判据全部落到真实读写站点与传输键，**不按命名猜**。
    """

    @staticmethod
    def _w(manifest_slice: dict) -> dict:
        return manifest_slice["writeoff_capability_resolution"]

    def test_verdict_is_binary_and_not_a_single_adjudication(self, manifest_slice: dict) -> None:
        w = self._w(manifest_slice)
        assert w["is_client_only"] is False
        assert w["verdict"] == "must_be_adapter_borne"
        assert "AC 12.8" in w["verdict_why"], (
            "必须写明「纯客户端本身不是裁 single 的理由，AC 12.8 的唯一判据是无 HTML 对端」"
        )
        for e in manifest_slice["independent_entries"]:
            assert e["capability"] != "single_onlyoffice"
            assert e["capability"] != "single_html"

    def test_owner_constants_are_read_from_source_not_copied(self, manifest_slice: dict) -> None:
        """第③边：impl 常量**现读**，不是抄一份快照。"""
        w = self._w(manifest_slice)
        owner = ROOT / w["owner_module"]
        for name, decl in w["owner_constants"].items():
            value, line = _const_literal(owner, name)
            assert value == decl["value"], (
                f"{name}: 声明 {decl['value']!r} != 现读 {value!r}"
            )
            assert decl["source_ref"] == (
                f"{owner.relative_to(ROOT).as_posix()}#L{line}"
            ), f"{name}: source_ref 行号与现读不符"

    def test_declared_transport_keys_hit_counts_recompute(
        self, manifest_slice: dict, k_files: list[pathlib.Path]
    ) -> None:
        w = self._w(manifest_slice)
        for key, expected in w["declared_key_hit_counts"].items():
            hits = _exact_literal_hits(k_files, key)
            assert len(hits) == expected, (
                f"键 {key!r} 的精确字面量命中数 {len(hits)} != 登记 {expected}"
                f"（前 5 处：{hits[:5]}）"
            )

    def test_guessed_keys_have_zero_hits(
        self, manifest_slice: dict, k_files: list[pathlib.Path]
    ) -> None:
        """非空反向分母：7 个「按命名规律该有」的键各 0 命中。"""
        w = self._w(manifest_slice)
        guessed = w["guessed_keys_with_zero_hits"]
        assert len(guessed) >= 5, "反向分母太小，判据接近空跑"
        for g in guessed:
            hits = _exact_literal_hits(k_files, g)
            assert hits == [], f"猜测键 {g!r} 竟有命中：{hits}"

    def test_bare_sheet_code_is_registered_as_not_a_transport_key(
        self, manifest_slice: dict, k_files: list[pathlib.Path]
    ) -> None:
        """反例登记：`K1-9` 有 12 处命中但它是 sheet 码，不得当传输键存在的证据。"""
        w = self._w(manifest_slice)
        for code, expected in w["bare_sheet_code_hits"].items():
            hits = _exact_literal_hits(k_files, code)
            assert len(hits) == expected, (
                f"sheet 码 {code!r} 命中数 {len(hits)} != 登记 {expected}"
            )
            assert code not in [t["key"] for t in w["transport_keys"]], (
                f"{code!r} 被误列进 transport_keys"
            )

    def test_write_path_hops_are_all_real(self, manifest_slice: dict) -> None:
        w = self._w(manifest_slice)
        assert len(w["write_path"]) == 5
        for hop in w["write_path"]:
            ref = hop["source_ref"]
            p = _resolve_repo(ref)
            assert p.exists(), f"hop {hop['hop']}: 文件不存在 {ref}"
            _line_at(ref)
        # hop 2/3/4/5 的关键代码形态必须真在那一带
        tab = _strip_comments(WRITEOFF_TAB.read_text(encoding="utf-8"))
        assert "buildSavePayload()" in tab and "emit('save'" in tab, (
            "K1TabWriteoffCheck 不再经 emit('save') 上抛 ⇒ 写路径断了，登记需重算"
        )
        host = _strip_comments(
            (WP_COMPONENTS / "GtK1OtherReceivables.vue").read_text(encoding="utf-8"))
        assert "persistence.save(" in host, "宿主不再调 persistence.save"
        cp = _strip_comments(CHECKLIST_PERSISTENCE.read_text(encoding="utf-8"))
        assert "/checklist-responses" in cp
        router = _strip_comments(CHECKLIST_ROUTER.read_text(encoding="utf-8"))
        assert re.search(r'@router\.put\(""', router)

    def test_property_24_derived_totals_are_computed_and_write_only(
        self, manifest_slice: dict
    ) -> None:
        """Property 24 的**真分母**：两个合计键是派生态且 HTML 侧只写不读。"""
        src = _strip_comments(WRITEOFF_OWNER.read_text(encoding="utf-8"))
        assert re.search(
            r"const reversalTotal:\s*ComputedRef<number>\s*=\s*computed", src
        ), "reversalTotal 不再是 computed ⇒ 「派生值」这个前提不成立"
        assert re.search(
            r"const writeoffTotal:\s*ComputedRef<number>\s*=\s*computed", src
        )
        # totals 数组由这两个 computed 产出
        m = re.search(r"const totals = \[(.*?)\]", src, re.S)
        assert m, "buildSavePayload 里找不到 totals 数组"
        body = m.group(1)
        assert "reversalTotal.value" in body and "writeoffTotal.value" in body
        assert "REVERSAL_TOTAL_KEY" in body and "WRITEOFF_TOTAL_KEY" in body
        # 反序列化入口只读 5 个字段，**不读**两个 total ⇒ 目前是受保护形态
        load = re.search(r"function load\(\): void \{(.*?)\n  \}", src, re.S)
        assert load, "找不到 load()"
        lbody = load.group(1)
        for field in ("tables", "auditProcedures", "auditNote", "conclusion", "conclusionOption"):
            assert field in lbody, f"load() 不再读 {field}"
        assert "REVERSAL_TOTAL_KEY" not in lbody and "WRITEOFF_TOTAL_KEY" not in lbody, (
            "🔴 load() 开始回读派生合计键 ⇒ summary 不再是「只写不读」的受保护形态，"
            "Property 24 的这条前提性质失守"
        )
        # 空分母那侧的前提
        for e in manifest_slice["independent_entries"]:
            assert e["published_representation"] is None

    def test_oo_counterpart_sheet_exists_but_is_unmapped(
        self, manifest_slice: dict, k1_template: pathlib.Path
    ) -> None:
        w = self._w(manifest_slice)
        assert w["oo_counterpart_status"] == "absent"
        names = _sheet_names(k1_template)
        assert any("K1-9" in n for n in names), (
            f"权威模板里没有 K1-9 sheet ⇒ oo_counterpart_status 的论证需重算（现有 {names}）"
        )
        for e in manifest_slice["independent_entries"]:
            if _wp_code_of(e) == "K1":
                assert e["adapter_id"] is None and e["definition_bundle"] is None


# ════════════════════════════════════════════════════════════════════════════
class TestAccountScopeAndFourTableReuse:
    """tasks.md 正文第一条：「复用 K1/K2 scope 与四表真源」。"""

    @staticmethod
    def _a(manifest_slice: dict) -> dict:
        return manifest_slice["account_scope_and_four_table_audit"]

    def test_k1_and_k2_scope_derive_from_render_issued_tb_source_codes(
        self, manifest_slice: dict
    ) -> None:
        """判据是「真从真源派生」，不是「抄第二份常量」。"""
        fe = self._a(manifest_slice)["frontend_scope_paradigm"]
        assert fe["verdict"] == "clean"
        for name in ("k1AccountScope.ts", "k2AccountScope.ts"):
            p = WP_COMPOSABLES / name
            src = _strip_comments(p.read_text(encoding="utf-8"))
            assert re.search(
                r"from '\./shared/tbSourceCodes'", src
            ), f"{name}: 未 import 共享的 tbSourceCodes ⇒ 不是从真源派生"
            assert "TbSourceCodes" in src
            assert re.search(r"src\?\.gross_standard", src), (
                f"{name}: 查询函数没有以 render 下发的 gross_standard 为一等入参"
            )
            assert "tbQueryCodes(" in src, (
                f"{name}: 没走共享的 tbQueryCodes（兜底只能作第二参）"
            )
        # 兜底常量只作兜底：K1 的坏账兜底必须是专属 1231-03 而不是宽口径 1231
        k1 = _strip_comments((WP_COMPOSABLES / "k1AccountScope.ts").read_text(encoding="utf-8"))
        m = re.search(r"K1_BAD_DEBT_FALLBACK_STANDARD = '([^']+)'", k1)
        assert m and m.group(1) == "1231-03", (
            f"K1 坏账兜底码变成 {m and m.group(1)!r} —— 宽口径 1231 会把 D1/D2 的坏账吃进来"
        )
        k2 = _strip_comments((WP_COMPOSABLES / "k2AccountScope.ts").read_text(encoding="utf-8"))
        m2 = re.search(r"K2_GROSS_FALLBACK_STANDARD = '([^']+)'", k2)
        assert m2 and m2.group(1) == "1901"
        m3 = re.search(r"K2_WRONG_LEGACY_ACCOUNT = '([^']+)'", k2)
        assert m3 and m3.group(1) == "1231", (
            "K2 的「曾被误用科目」守卫常量不见了 —— 那条反向判据会失效"
        )

    def test_every_entry_has_its_own_per_cycle_scope_file(self, manifest_slice: dict) -> None:
        declared = self._a(manifest_slice)["frontend_scope_paradigm"]["per_entry_scope_files"]
        assert len(declared) == 13
        for rel in declared:
            assert (ROOT / rel).exists(), f"per-cycle 科目真源不存在：{rel}"

    def test_backend_k_cycle_specs_uses_the_shared_report_line_spec(
        self, manifest_slice: dict
    ) -> None:
        be = self._a(manifest_slice)["backend_four_table_reuse"]
        assert be["verdict"] == "clean"
        src = K_CYCLE_SPECS_PY.read_text(encoding="utf-8")
        assert "from .report_line_accounts import ReportLineAccountSpec" in src, (
            "k_cycle_specs 不再复用共享的 ReportLineAccountSpec ⇒ 抄了第二份科目定位逻辑"
        )
        assert "-> ReportLineAccountSpec" in src

    @staticmethod
    def _main_strategies() -> list[pathlib.Path]:
        """🔴 主渲染策略按「是否 import k_cycle_specs」筛，**不按文件名前缀**。

        `glob('_k*.py')` 会捞到 42 个文件：13 个主策略 + 3 个 K0 函证族 + 26 个
        `_ai_generate` / `_import_export` 辅助策略。按前缀判会把 `_k0_confirmation.py`
        当成损益侧 entry 而假红（首轮实测）。
        """
        return sorted(
            p for p in RENDER_STRATEGY_DIR.glob("_k*.py")
            if "app.services.four_table.k_cycle_specs" in p.read_text(encoding="utf-8")
        )

    def test_all_thirteen_render_strategies_import_k_cycle_specs(
        self, manifest_slice: dict
    ) -> None:
        be = self._a(manifest_slice)["backend_four_table_reuse"]
        found = [p.relative_to(ROOT).as_posix() for p in self._main_strategies()]
        assert len(found) == 13, f"只有 {len(found)} 个 K render 策略复用声明真源：{found}"
        assert found == be["main_render_strategies"], "主策略清单登记与现算不符"
        assert manifest_slice["honest_adjudication_summary"][
            "backend_render_strategies_importing_k_cycle_specs"] == 13
        assert be["all_k_prefixed_strategy_files"] == len(
            list(RENDER_STRATEGY_DIR.glob("_k*.py"))), (
            "`_k*.py` 文件总数登记与磁盘不符 —— 「按前缀筛会捞到多少」这个论证会失效"
        )
        assert be["all_k_prefixed_strategy_files"] > 13, (
            "🔴 前缀口径与真口径必须真的不同，否则 main_render_strategy_selector 的告警是空的"
        )
        k0 = sorted(p.relative_to(ROOT).as_posix() for p in RENDER_STRATEGY_DIR.glob("_k0_*.py"))
        assert k0 == be["k0_render_strategies"] and k0, (
            "K0 函证族策略清单不符 —— 它是 K0 排除结论的第四条实证"
        )
        for p in k0:
            assert "k_cycle_specs" not in (ROOT / p).read_text(encoding="utf-8"), (
                f"{p} 竟 import 了 k_cycle_specs ⇒ K0 走函证族这个结论要重算"
            )

    def test_balance_side_uses_leaf_aggregation_and_pl_side_uses_pl_render(
        self, manifest_slice: dict
    ) -> None:
        """🔴 分支判据：统一要求 select_leaves 会把 6 条损益 entry 假红。"""
        be = self._a(manifest_slice)["backend_four_table_reuse"]
        bs = set(be["balance_side_entries"])
        pl = set(be["pl_side_entries"])
        assert bs & pl == set() and len(bs | pl) == 13
        seen: set[str] = set()
        for p in self._main_strategies():
            m = re.match(r"_k(\d+)_", p.name)
            assert m, p.name
            code = f"K{m.group(1)}"
            seen.add(code)
            src = p.read_text(encoding="utf-8")
            if code in bs:
                assert "select_leaves" in src or "aggregate_leaves" in src, (
                    f"{p.name}: 资产负债侧却没走共享叶子聚合"
                )
            else:
                assert code in pl, f"{p.name}: {code} 既不在资产负债侧也不在损益侧"
                assert "render_pl_cycle" in src, (
                    f"{p.name}: 损益侧却没走共享 pl_render.render_pl_cycle"
                )
                assert "select_leaves" not in src, (
                    f"{p.name}: 损益侧不该直接用 select_leaves（它在 pl_occurrence 里）"
                )
        assert seen == bs | pl, f"主策略覆盖的 entry 集合不全：漏 {sorted((bs | pl) - seen)}"

    def test_no_duplicated_account_location_or_aggregation_in_k_domain(
        self, manifest_slice: dict
    ) -> None:
        """反向：K 专属文件里不得自己实现科目前缀匹配 / 方向带符号求和。"""
        be = self._a(manifest_slice)["backend_four_table_reuse"]
        assert be["duplicate_aggregation_found"] == 0
        suspects: list[str] = []
        dup = re.compile(r"account_code\s*\.\s*like|closing_direction\s*==")
        for p in sorted(RENDER_STRATEGY_DIR.glob("_k*.py")):
            src = p.read_text(encoding="utf-8")
            if dup.search(src) and "four_table" not in src:
                suspects.append(p.name)
        assert suspects == [], f"这些 K 文件自己写了科目定位/方向聚合：{suspects}"

    def test_cross_lock_carrier_exists_and_is_not_duplicated_here(
        self, manifest_slice: dict
    ) -> None:
        fe = self._a(manifest_slice)["frontend_scope_paradigm"]
        rel = fe["cross_lock_carrier"].split(" ")[0].strip()
        p = ROOT / rel
        assert p.exists(), (
            f"声明的交叉锁守卫不存在：{rel} —— 「不再造第二份」的前提是那一份真在"
        )
        src = p.read_text(encoding="utf-8")
        assert "k_cycle_specs" in src, "交叉锁守卫不再直读后端 py 源码"


# ════════════════════════════════════════════════════════════════════════════
class TestProperty23StaticStructure:
    @staticmethod
    def _pii(manifest_slice: dict) -> dict:
        return manifest_slice["dynamic_row_identity"]["positional_identity_inventory"]

    def test_positional_identity_inventory_is_exhaustive_and_partitioned(
        self, manifest_slice: dict, k_files: list[pathlib.Path]
    ) -> None:
        pii = self._pii(manifest_slice)
        hits = _positional_identity_hits(k_files)
        assert len(hits) == pii["total_hits"] == 48, (
            f"位置化命中总数 {len(hits)} != 登记 {pii['total_hits']}"
        )
        fams: dict[str, list[str]] = {"a": [], "b": [], "c": []}
        for ref, _key, val in hits:
            fams[_family_of(val)].append(ref)
        assert sorted(fams["a"]) == sorted(pii["family_a_pure_ordinal"]["hits"])
        assert sorted(fams["b"]) == sorted(pii["family_b_index_as_fallback"]["hits"])
        assert sorted(fams["c"]) == sorted(
            pii["family_c_generated_opaque_must_not_be_flagged"]["hits"])
        assert len(fams["a"]) == pii["family_a_pure_ordinal"]["count"] == 32
        assert len(fams["b"]) == pii["family_b_index_as_fallback"]["count"] == 13
        assert len(fams["c"]) == pii[
            "family_c_generated_opaque_must_not_be_flagged"]["count"] == 3
        assert len(fams["a"]) + len(fams["b"]) + len(fams["c"]) == pii["total_hits"]
        # 三族互斥
        assert set(fams["a"]) & set(fams["b"]) == set()
        assert set(fams["b"]) & set(fams["c"]) == set()

    def test_defect_count_and_per_entry_distribution_recompute(
        self, manifest_slice: dict, k_files: list[pathlib.Path]
    ) -> None:
        pii = self._pii(manifest_slice)
        hits = _positional_identity_hits(k_files)
        defects = [ref for ref, _k, v in hits if _family_of(v) in ("a", "b")]
        assert len(defects) == pii["defect_hits_total"] == 45

        def owner(ref: str) -> str | None:
            for rx in (r"/workpaper/k(\d+)/", r"/(?:use)?[kK](\d+)[A-Z]", r"/GtK(\d+)"):
                m = re.search(rx, ref)
                if m:
                    return "K" + m.group(1)
            return None

        counted: dict[str, int] = {}
        for ref in defects:
            o = owner(ref)
            assert o, f"无法归属的命中：{ref}"
            counted[o] = counted.get(o, 0) + 1
        assert counted == pii["defect_hits_by_entry"], (
            f"逐 entry 分布不符：现算 {counted} != 登记 {pii['defect_hits_by_entry']}"
        )
        assert sum(counted.values()) == 45

    def test_entries_with_zero_defect_hits_are_really_zero(
        self, manifest_slice: dict, k_files: list[pathlib.Path]
    ) -> None:
        """🔴「某族 0 命中」也要落成判据 —— K2/K4/K10/K13 将来新增一处会立刻打红。"""
        pii = self._pii(manifest_slice)
        zero = pii["entries_with_zero_defect_hits"]
        assert zero == ["K2", "K4", "K10", "K13"], f"零命中集合变了：{zero}"
        assert set(zero) & set(pii["defect_hits_by_entry"]) == set()
        assert len(zero) + len(pii["defect_hits_by_entry"]) == 13
        hits = _positional_identity_hits(k_files)
        for code in zero:
            n = int(code[1:])
            bad = [
                ref for ref, _k, v in hits
                if _family_of(v) in ("a", "b")
                and (f"/workpaper/k{n}/" in ref or re.search(r"/(?:use)?[kK]%d[A-Z]" % n, ref))
            ]
            assert bad == [], f"{code} 声称 0 处位置化缺陷，现算却有 {bad}"

    def test_display_sequence_sites_are_not_flagged(
        self, manifest_slice: dict, k_files: list[pathlib.Path]
    ) -> None:
        pii = self._pii(manifest_slice)
        fam_d = pii["family_d_display_ordinal_must_not_be_flagged"]
        sites = _display_seq_sites(k_files)
        assert len(sites) == fam_d["count"] == 38
        assert sorted(sites) == sorted(fam_d["hits"])
        # family_d 不计入 total_hits
        hits = {ref for ref, _k, _v in _positional_identity_hits(k_files)}
        assert fam_d["is_defect"] is False
        assert len(set(sites) - hits) >= 30, (
            "🔴 展示序号站点几乎全被身份键正则命中了 ⇒ 两个口径没分开，family_d 的排除失效"
        )

    def test_hardcoded_scan_recomputes_including_the_non_zero_one(
        self, manifest_slice: dict, k_files: list[pathlib.Path]
    ) -> None:
        """🔴 不得照抄 J 的「六个全 0」—— K 的第六个是 13。"""
        hs = manifest_slice["dynamic_row_identity"]["hardcoded_scan_result"]
        assert set(hs["patterns"]) == set(_HARDCODED_PATTERNS)
        for name, expected in hs["patterns"].items():
            hits = _hardcoded_hits(k_files, name)
            assert len(hits) == expected, (
                f"模式 {name} 命中数 {len(hits)} != 登记 {expected}（前 5：{hits[:5]}）"
            )
            assert sorted(hits) == sorted(hs["hits"][name])
        zeros = [k for k, v in hs["patterns"].items() if v == 0]
        assert len(zeros) == 5, f"应有 5 个模式为 0，实得 {zeros}"
        assert hs["patterns"]["positional_row_id_template"] == 13
        assert hs["why_the_sixth_is_not_zero"].strip()
        assert hs["not_claimed_as_evidence_of_quality"].strip()

    def test_the_non_zero_pattern_is_a_subset_of_the_positional_hits(
        self, manifest_slice: dict, k_files: list[pathlib.Path]
    ) -> None:
        """两个扫描口径互相印证，而不是各说各话。"""
        hs = manifest_slice["dynamic_row_identity"]["hardcoded_scan_result"]
        tmpl = set(hs["hits"]["positional_row_id_template"])
        pos = {ref for ref, _k, _v in _positional_identity_hits(k_files)}
        assert tmpl <= pos, (
            f"positional_row_id_template 有不在身份键命中里的行：{sorted(tmpl - pos)}"
        )
        assert len(tmpl) == 13

    def test_declared_dynamic_tables_avoid_forbidden_identity_kinds(
        self, manifest_slice: dict
    ) -> None:
        dri = manifest_slice["dynamic_row_identity"]
        forbidden = set(dri["forbidden_identity_kinds"])
        assert forbidden
        for t in dri["tables"]:
            ident = t["row_identity"]
            assert ident["kind"] not in forbidden, (
                f"{t['table_key']}: kind={ident['kind']} 落在 forbidden_identity_kinds 里"
            )
            ref = ident["source_ref"]
            assert _resolve_repo(ref).exists()
            line = _line_at(ref)
            assert ident["identity_field"] in line, (
                f"{t['table_key']}: source_ref 那一行不含身份字段 {ident['identity_field']}"
                f"（实际：{line.strip()!r}）"
            )
            assert _ENTROPY.search(line), (
                f"{t['table_key']}: 声称 generated_opaque_string，那一行却没有熵来源"
            )

    def test_the_two_containers_partition_all_hits(self, manifest_slice: dict) -> None:
        """tables[] 与 positional_identity_inventory 的并集覆盖全部、交集为空。"""
        dri = manifest_slice["dynamic_row_identity"]
        pii = dri["positional_identity_inventory"]
        table_refs = {t["row_identity"]["source_ref"] for t in dri["tables"]}
        defect_refs = set(pii["family_a_pure_ordinal"]["hits"]) | set(
            pii["family_b_index_as_fallback"]["hits"])
        assert table_refs & defect_refs == set(), (
            f"tables[] 里出现了已登记为缺陷的行：{sorted(table_refs & defect_refs)}"
        )
        fam_c = set(pii["family_c_generated_opaque_must_not_be_flagged"]["hits"])
        assert table_refs <= fam_c, (
            f"tables[] 的主身份必须落在 family_c（有熵）里：{sorted(table_refs - fam_c)}"
        )
        assert pii["registered_as"] == "BP-8"


# ════════════════════════════════════════════════════════════════════════════
class TestProperty3And20:
    def test_no_slice_entry_has_a_contract(self, manifest_slice: dict) -> None:
        k_ids = {e["entry_id"] for e in manifest_slice["independent_entries"]}
        for p in sorted(CONTRACT_DIR.glob("*.json")):
            doc = _load(p)
            assert (doc.get("review") or {}).get("entry_id") not in k_ids

    def test_no_slice_entry_has_a_registered_adapter(self, manifest_slice: dict) -> None:
        reg = REGISTRY.read_text(encoding="utf-8")
        for e in manifest_slice["independent_entries"]:
            assert e["adapter_id"] is None
            assert e["entry_id"] not in reg, (
                f"{e['entry_id']} 出现在 adapter registry 里 ⇒ adapter_id 为 null 是假的"
            )

    def test_no_entry_claims_bidirectional(self, manifest_slice: dict) -> None:
        s = manifest_slice["honest_adjudication_summary"]
        assert s["adjudicated_as_bidirectional"] == 0
        assert s["slice_counters"]["bidirectional_unverified"] == 0
        assert s["slice_counters"]["fake_bidirectional_claimed_verified"] == 0

    def test_host_templates_make_no_bidirectional_claim(self, manifest_slice: dict) -> None:
        for e in manifest_slice["independent_entries"]:
            tpl = _strip_comments(
                _vue_template((ROOT / e["host_path"]).read_text(encoding="utf-8")))
            for claim in ("可双向回写", "双向回写"):
                assert claim not in tpl, (
                    f"{e['entry_id']}: 宿主模板出现「{claim}」字样，但 adapter 未注册（AC 1.4）"
                )

    def test_ac14_notice_single_source_exists_and_is_not_duplicated(self) -> None:
        """AC 1.4 文案真源唯一；不得有第二份实现。"""
        assert NOTICE_MODULE.exists() and NOTICE_COMPONENT.exists()
        comp = NOTICE_COMPONENT.read_text(encoding="utf-8")
        assert "workpaperEntrySyncNotice" in comp, (
            "提示组件不再读单一真源模块 ⇒ 文案抄了第二份"
        )
        dupes = [
            p.relative_to(ROOT).as_posix() for p in _all_frontend()
            if p != NOTICE_MODULE and p != NOTICE_COMPONENT and not _is_test_path(p)
            and re.search(r"const\s+\w*SYNC_NOTICE\w*\s*[:=]", _cached_text(p))
        ]
        assert dupes == [], f"AC 1.4 文案出现第二份真源：{dupes}"

    def test_bp7_is_registered_because_no_host_mounts_the_notice(
        self, manifest_slice: dict
    ) -> None:
        """AC 1.4 的义务未兑现 ⇒ 必须登记，且计数为 0 是现算的。"""
        mounted = 0
        for e in manifest_slice["independent_entries"]:
            src = _strip_comments((ROOT / e["host_path"]).read_text(encoding="utf-8"))
            has = NOTICE_COMPONENT_NAME in src or "workpaperEntrySyncNotice" in src
            assert e["ui_toolbar_gate"]["mounts_ac14_notice"] == has, (
                f"{e['entry_id']}: mounts_ac14_notice 登记与现算不符"
            )
            mounted += int(has)
        assert mounted == 0, (
            f"🔴 有 {mounted} 个宿主挂上了 AC 1.4 提示 ⇒ BP-7 可解除，登记必须更新"
        )
        assert manifest_slice["honest_adjudication_summary"][
            "entries_mounting_the_ac14_notice"] == 0
        assert any(b["id"] == "BP-7" for b in manifest_slice["blocking_preconditions"])

    def test_toolbar_gate_anchor_is_unique_resolvable_and_contains_the_switcher(
        self, manifest_slice: dict
    ) -> None:
        """AC 1.4 的挂载点必须落在该 entry 自己声明的 ui_toolbar_gate 区块内。"""
        for e in manifest_slice["independent_entries"]:
            gate = e["ui_toolbar_gate"]
            ref = gate["anchor"]
            assert _resolve_repo(ref) == ROOT / e["host_path"]
            code = _wp_code_of(e)
            line = _line_at(ref)
            assert f'class="{code.lower()}-header-toolbar"' in line, (
                f"{e['entry_id']}: 锚点行不是 toolbar 容器（实际 {line.strip()!r}）"
            )
            src = (ROOT / e["host_path"]).read_text(encoding="utf-8")
            anchors = [
                i for i, l in enumerate(_strip_comments(src).split("\n"), 1)
                if f'class="{code.lower()}-header-toolbar"' in l
            ]
            assert anchors == [_line_no_of(ref)], (
                f"{e['entry_id']}: toolbar 锚点不唯一（{anchors}）⇒ 区块截取会有歧义"
            )
            block = _toolbar_block(_strip_comments(_vue_template(src)), _line_no_of(ref))
            assert "el-segmented" in block, (
                f"{e['entry_id']}: 声明的 toolbar 区块里没有模式切换器 ⇒ 锚点选错了"
            )
            assert gate["segmented_sites_in_gate"], "区块内分段器行号未登记"
            assert set(gate["segmented_sites_in_gate"]) <= set(
                gate["segmented_sites_before_comment_strip"]), (
                f"{e['entry_id']}: 剥注释后的分段器行号不是剥注释前的子集 ⇒ 口径矛盾"
            )
            stripped_sites = [
                i for i, l in enumerate(_strip_comments(src).split("\n"), 1)
                if "el-segmented" in l
            ]
            assert stripped_sites == gate["segmented_sites_in_gate"], (
                f"{e['entry_id']}: segmented_sites_in_gate 登记与现算不符"
            )

    def test_property_20_denominator_declares_what_is_not_claimed(
        self, manifest_slice: dict
    ) -> None:
        pd = manifest_slice["property_denominators"]["property_20"]
        assert pd["not_claimed_passing"] is True
        assert pd["not_claimed_passing_part"].strip()
        assert pd["carrier_for_the_empty_part"].strip()
        assert "分母为空" in pd["k_cycle_denominator"]

    def test_property_24_denominator_declares_what_is_not_claimed(
        self, manifest_slice: dict
    ) -> None:
        pd = manifest_slice["property_denominators"]["property_24"]
        assert pd["not_claimed_passing"] is True
        assert pd["not_claimed_passing_part"].strip()
        assert pd["carrier_for_the_empty_part"].strip()
        assert "6.6" in manifest_slice["requirements_covered"] or "6.6" in pd["what_it_asserts"]


# ════════════════════════════════════════════════════════════════════════════
class TestProperty28DefinitionDriftFailClosed:
    def test_authoritative_template_digests_recompute(self, manifest_slice: dict) -> None:
        for f in manifest_slice["authoritative_templates"]["files"]:
            p = K_TEMPLATE_DIR / f["name"]
            assert p.exists(), f"权威模板不存在：{p}"
            assert p.stat().st_size == f["size"], f"{f['name']}: size 漂移"
            assert _sha256_of(p) == f["sha256"], f"{f['name']}: sha256 漂移（fail closed）"

    def test_registered_file_set_equals_the_disk_set(self, manifest_slice: dict) -> None:
        on_disk = {n for n in os.listdir(K_TEMPLATE_DIR) if not n.startswith("~$")}
        declared = {f["name"] for f in manifest_slice["authoritative_templates"]["files"]}
        assert declared == on_disk, (
            f"多写 {sorted(declared - on_disk)}，漏写 {sorted(on_disk - declared)}"
        )
        assert len(on_disk) == 14

    def test_sheet_inventory_recomputes_and_names_are_not_stripped(
        self, manifest_slice: dict
    ) -> None:
        total = 0
        for f in manifest_slice["authoritative_templates"]["files"]:
            names = _sheet_names(K_TEMPLATE_DIR / f["name"])
            assert names == f["sheets"], f"{f['name']}: sheet 清单漂移"
            assert len(names) == f["sheet_count"]
            total += len(names)
            for n in names:
                assert n == n.strip() or True  # 记录用；真判据在下一行
        assert total == manifest_slice["honest_adjudication_summary"][
            "authoritative_template_sheets_total"]
        stripped = [
            (f["name"], n) for f in manifest_slice["authoritative_templates"]["files"]
            for n in f["sheets"] if n != n.strip()
        ]
        assert stripped == [], (
            f"🔴 出现了带首尾空白的 sheet 名：{stripped} —— sheet_name_policy 的「0 命中」失效，"
            "所有按名取 sheet 的判据都要改成不 strip"
        )

    def test_lock_files_are_really_absent(self, manifest_slice: dict) -> None:
        locks = [n for n in os.listdir(K_TEMPLATE_DIR) if n.startswith("~$")]
        assert locks == [], f"出现 Office 锁文件 {locks} —— 枚举口径要跳过它们"
        assert "~$" in manifest_slice["authoritative_templates"]["lock_file_policy"]

    def test_reference_copy_status_is_honest(self, manifest_slice: dict) -> None:
        at = manifest_slice["authoritative_templates"]
        assert at["reference_copy_status"] == "absent_in_worktree"
        assert not (ROOT / "基础数据").exists(), (
            "参考副本目录出现了 ⇒ reference_copy_status 必须改，并做双处比对"
        )
        ref_dir = TEMPLATE_DIR / "_reference"
        assert ref_dir.is_dir(), "_reference 目录不见了 —— 那条防误认的说明需要重算"
        assert "_reference" in at["reference_copy_note"], (
            "必须写明 `backend/wp_templates/_reference/` 不是模板库副本"
        )
        assert len(list(ref_dir.iterdir())) == 6

    def test_template_owner_mapping_is_injective_with_reasoned_nulls(
        self, manifest_slice: dict
    ) -> None:
        entry_ids = {e["entry_id"] for e in manifest_slice["independent_entries"]}
        owners: list[str] = []
        for f in manifest_slice["authoritative_templates"]["files"]:
            o = f["belongs_to_entry"]
            if o is None:
                assert f.get("excluded_reason"), f"{f['name']}: null 却无 excluded_reason（SR-8）"
            else:
                assert o in entry_ids, f"{f['name']}: belongs_to_entry={o} 不命中本 slice（SR-8）"
                owners.append(o)
        assert len(owners) == len(set(owners)) == 13, "模板→entry 不是单射"

    def test_template_resolution_audit_recomputes(self, manifest_slice: dict) -> None:
        sys.path.insert(0, str(BACKEND))
        try:
            from app.services.wp_template_finder import (  # type: ignore
                find_template_file, find_template_file_any,
            )
        finally:
            pass
        audit = manifest_slice["template_resolution_audit"]
        assert audit["result"] == "clean"
        measured = audit["measured"]
        assert len(measured) == audit["measured_code_count"] == 168
        fallbacks = 0
        for rec in measured:
            a = find_template_file(rec["code"])
            b = find_template_file_any(rec["code"])
            an = pathlib.Path(str(a)).name if a else None
            bn = pathlib.Path(str(b)).name if b else None
            assert an == rec["find"], f"{rec['code']}: find_template_file 结果漂移 {an!r}"
            assert bn == rec["find_any"], f"{rec['code']}: find_template_file_any 漂移 {bn!r}"
            assert an == bn, f"{rec['code']}: 两个解析函数不一致 —— audit 的 clean 结论失效"
            if an:
                m = re.match(r"^K(\d+)", rec["code"])
                if m and not an.startswith(f"K{m.group(1)} "):
                    fallbacks += 1
        assert fallbacks == audit["cross_workbook_fallback_hits"] == 0
        assert audit["two_functions_agree"] is True

    def test_program_table_codes_resolve_to_none_but_k0_does_not(
        self, manifest_slice: dict
    ) -> None:
        """🔴 不得照抄 J 的「X0 返回 None」—— K0 是真册。"""
        measured = {r["code"]: r for r in manifest_slice["template_resolution_audit"]["measured"]}
        for n in range(0, 14):
            rec = measured[f"K{n}A"]
            assert rec["find"] is None and rec["find_any"] is None, (
                f"K{n}A 程序表码竟解析到 {rec} —— 程序表是 sheet 不是册"
            )
        k0 = measured["K0"]
        assert k0["find"] == k0["find_any"] == "K0 管理循环函证.xlsx", (
            f"K0 应解析到真实函证册（与 J0 相反），实得 {k0}"
        )
        assert "J0" in manifest_slice["template_resolution_audit"]["k0_code_note"]

    def test_bundle_layer_drift_is_not_claimed(self, manifest_slice: dict) -> None:
        p = manifest_slice["property_denominators"]["property_3_21_22_23_28_not_claimed"]
        assert p["not_claimed_passing"] is True
        blob = json.dumps(p, ensure_ascii=False)
        assert "bundle" in blob, "必须写明 bundle 层 drift 分母为空、不宣称通过"


# ════════════════════════════════════════════════════════════════════════════
class TestProperty69EvidenceAndCounters:
    def test_unverifiable_entries_carry_non_empty_reasons(self, manifest_slice: dict) -> None:
        n = 0
        for e in manifest_slice["independent_entries"]:
            ev = e["evidence"]
            if ev["verification_state"] == "UNVERIFIABLE":
                n += 1
                assert isinstance(ev["unverifiable_reasons"], list) and ev["unverifiable_reasons"], (
                    f"{e['entry_id']}: SR-7 违反"
                )
            else:
                assert ev["sync_test_run_id"] and ev["required_scenario_set_digest"], (
                    f"{e['entry_id']}: 非 UNVERIFIABLE 却没有 run id / digest ⇒ 空口声称"
                )
            n_bp = {r.split("：")[0] for r in ev.get("unverifiable_reasons") or []}
            declared = {b["id"] for b in manifest_slice["blocking_preconditions"]}
            assert n_bp <= declared | {""}, f"{e['entry_id']}: reasons 引用了不存在的 BP {n_bp}"
        assert n == manifest_slice["honest_adjudication_summary"]["entries_left_unverifiable"] == 13

    def test_contract_test_points_at_this_guard(self, manifest_slice: dict) -> None:
        me = _THIS.relative_to(ROOT).as_posix()
        for e in manifest_slice["independent_entries"]:
            assert e["evidence"]["contract_test"] == me, (
                f"{e['entry_id']}: contract_test 应指向本守卫（{me}）"
            )

    def test_summary_counters_recompute_from_the_entries(self, manifest_slice: dict) -> None:
        E = manifest_slice["independent_entries"]
        s = manifest_slice["honest_adjudication_summary"]
        assert s["total_independent"] == len(E) == 13
        assert s["capability_verdict_pending"] == sum(1 for e in E if e["capability"] is None)
        for cap in CAPABILITY_ENUM:
            assert s[f"adjudicated_as_{cap}"] == sum(
                1 for e in E if e["capability"] == cap)
        assert s["html_counterpart_exists"] == sum(
            1 for e in E if e["html_counterpart_verdict"] == "exists")
        assert s["html_counterpart_none"] == sum(
            1 for e in E if e["html_counterpart_verdict"] == "none")
        assert s["html_counterpart_exists"] + s["html_counterpart_none"] == len(E)
        assert s["reachable_hosts_in_cycle"] == len({e["host_path"] for e in E})
        assert s["hosts_with_oo_mount"] == sum(1 for e in E if e["mount_count"] > 0)
        assert s["entries_mounting_the_ac14_notice"] == sum(
            1 for e in E if e["ui_toolbar_gate"]["mounts_ac14_notice"])
        assert s["finalized_published_representations"] == sum(
            1 for e in E if e["published_representation"] is not None)
        assert s["blocking_preconditions_total"] == len(manifest_slice["blocking_preconditions"])

    def test_derived_counters_recompute_from_their_own_sections(
        self, manifest_slice: dict
    ) -> None:
        s = manifest_slice["honest_adjudication_summary"]
        tpl = manifest_slice["authoritative_templates"]["files"]
        assert s["authoritative_template_files"] == len(tpl) == 14
        assert s["template_files_with_owner_entry"] == sum(
            1 for f in tpl if f["belongs_to_entry"])
        assert s["template_files_without_owner_entry"] == sum(
            1 for f in tpl if not f["belongs_to_entry"])
        assert (s["template_files_with_owner_entry"]
                + s["template_files_without_owner_entry"] == 14)
        assert s["authoritative_template_sheets_total"] == sum(f["sheet_count"] for f in tpl)

        inv = manifest_slice["orphan_dual_mode_inventory"]
        assert s["orphan_dual_mode_modules"] == len(inv["modules"]) == 7
        assert s["live_dual_mode_modules"] == len(inv["live_modules"]) == 6
        assert s["orphan_dual_mode_second_order"] == 0 and s["orphan_barrels"] == 0
        assert s["shared_base_statement_consumers"] == SHARED_BASE_CONSUMERS
        assert s["shared_base_k_cycle_consumers"] == 0

        carriers = [e["dual_mode_carrier"]["kind"] for e in manifest_slice["independent_entries"]]
        assert s["hosts_with_inlined_iife_dual_mode"] == carriers.count("host_inlined_iife") == 7
        assert s["hosts_with_dedicated_composable_dual_mode"] == carriers.count(
            "dedicated_composable") == 6

        w = manifest_slice["writeoff_capability_resolution"]
        assert s["writeoff_transport_keys"] == len(w["transport_keys"]) == 3
        assert s["writeoff_is_client_only"] == w["is_client_only"] is False
        assert s["writeoff_write_path_hops"] == len(w["write_path"])
        assert s["guessed_transport_keys_with_zero_hits"] == len(w["guessed_keys_with_zero_hits"])

        a = manifest_slice["account_scope_and_four_table_audit"]
        assert s["per_entry_account_scope_files"] == len(
            a["frontend_scope_paradigm"]["per_entry_scope_files"]) == 13
        assert s["duplicate_account_aggregation_found"] == a[
            "backend_four_table_reuse"]["duplicate_aggregation_found"] == 0

        dri = manifest_slice["dynamic_row_identity"]
        pii = dri["positional_identity_inventory"]
        assert s["dynamic_row_identity_tables"] == len(dri["tables"])
        assert s["positional_identity_hits_total"] == pii["total_hits"] == 48
        assert s["positional_identity_defect_hits"] == pii["defect_hits_total"] == 45
        assert s["positional_identity_family_a"] == pii["family_a_pure_ordinal"]["count"]
        assert s["positional_identity_family_b"] == pii["family_b_index_as_fallback"]["count"]
        assert s["positional_identity_family_c"] == pii[
            "family_c_generated_opaque_must_not_be_flagged"]["count"]
        assert s["positional_identity_family_d"] == pii[
            "family_d_display_ordinal_must_not_be_flagged"]["count"]
        assert (s["positional_identity_family_a"] + s["positional_identity_family_b"]
                + s["positional_identity_family_c"] == s["positional_identity_hits_total"])
        assert s["entries_with_positional_identity_defects"] == len(pii["defect_hits_by_entry"])
        assert s["entries_without_positional_identity_defects"] == len(
            pii["entries_with_zero_defect_hits"])
        assert (s["entries_with_positional_identity_defects"]
                + s["entries_without_positional_identity_defects"] == 13)

        hs = dri["hardcoded_scan_result"]
        assert s["hardcoded_pattern_hits_total"] == sum(hs["patterns"].values()) == 13
        assert s["hardcoded_patterns_at_zero"] == sum(
            1 for v in hs["patterns"].values() if v == 0) == 5

        audit = manifest_slice["template_resolution_audit"]
        assert s["template_resolution_codes_measured"] == audit["measured_code_count"]
        assert s["template_resolution_cross_workbook_fallbacks"] == audit[
            "cross_workbook_fallback_hits"]

    def test_slice_counters_are_all_zero_except_unadjudicated(self, manifest_slice: dict) -> None:
        c = manifest_slice["honest_adjudication_summary"]["slice_counters"]
        assert c["unadjudicated"] == 13
        assert c["fake_bidirectional_claimed_verified"] == 0
        assert c["bidirectional_unverified"] == 0
        assert c["stale_evidence"] == 0
        assert c["note"].strip()

    def test_counting_notes_cover_every_nontrivial_counter(self, manifest_slice: dict) -> None:
        s = manifest_slice["honest_adjudication_summary"]
        notes = s["counting_notes"]
        assert notes, "counting_notes 为空 ⇒ 计数没有配方，抄错也看不出"
        for key in ("total_independent", "capability_verdict_pending",
                    "authoritative_template_files", "orphan_dual_mode_modules",
                    "positional_identity_hits_total", "hardcoded_pattern_hits_total",
                    "shared_base_k_cycle_consumers"):
            assert key in notes and notes[key].strip(), f"计数 {key} 没有 counting_note"
        assert "不是 0" in notes["hardcoded_pattern_hits_total"], (
            "必须显式提醒「不得照抄 J 的六个全 0」"
        )


# ════════════════════════════════════════════════════════════════════════════
class TestProperty70CrossEntryIsolation:
    def test_eight_slices_are_pairwise_disjoint(self, manifest_slice: dict) -> None:
        sets = {"K": {e["entry_id"] for e in manifest_slice["independent_entries"]}}
        for name, p in SIBLING_SLICE_PATHS.items():
            sets[name] = {e["entry_id"] for e in _load(p)["independent_entries"]}
        names = sorted(sets)
        pairs = 0
        for i, a in enumerate(names):
            for b in names[i + 1:]:
                pairs += 1
                assert sets[a] & sets[b] == set(), (
                    f"{a} 与 {b} 的 entry 集合相交：{sorted(sets[a] & sets[b])}"
                )
        assert pairs == 28, f"应算 28 个配对，实得 {pairs}"
        assert len(sets["K"]) == 13

    def test_sibling_slices_declared_match_the_disk(self, manifest_slice: dict) -> None:
        declared = set(manifest_slice["sibling_slices"])
        actual = {p.relative_to(ROOT).as_posix() for p in SIBLING_SLICE_PATHS.values()}
        assert declared == actual, f"sibling_slices 登记与磁盘不符：{declared ^ actual}"

    def test_cross_entry_isolation_assertions_are_present(self, manifest_slice: dict) -> None:
        iso = manifest_slice["cross_entry_isolation"]
        assert iso["rule"].strip()
        assert isinstance(iso["assertions"], list) and len(iso["assertions"]) >= 6

    def test_deletion_plan_paths_are_globally_unique_and_disjoint(
        self, deletion_plan: dict
    ) -> None:
        mine: set[str] = set()
        for m in deletion_plan["orphan_dual_mode_to_delete"]["modules"]:
            mine.add(m["file"])
        for m in deletion_plan["live_dual_mode_to_rewire_then_delete"]["modules"]:
            mine.add(m["file"])
        for e in deletion_plan["entries"]:
            mine |= set(e["legacy_composables_to_delete"])
        assert len(mine) == 13, f"删除路径集合应恰 13 个文件，实得 {sorted(mine)}"
        for name, p in SIBLING_SLICE_PATHS.items():
            plan = DATA / f"workpaper_sync_{name.lower()}_cycle_deletion_plan.json"
            if not plan.exists():
                continue
            other = set(re.findall(r'"([^"]*use\w+\.ts)"', plan.read_text(encoding="utf-8")))
            overlap = mine & other
            assert overlap == set(), f"与 {name} 循环的删除路径相交：{sorted(overlap)}"


# ════════════════════════════════════════════════════════════════════════════
class TestDeletionPlanConsistency:
    def test_plan_entries_mirror_the_slice_adjudication(
        self, deletion_plan: dict, manifest_slice: dict
    ) -> None:
        by_id = {e["entry_id"]: e for e in manifest_slice["independent_entries"]}
        assert len(deletion_plan["entries"]) == len(by_id) == 13
        for x in deletion_plan["entries"]:
            e = by_id[x["entry_id"]]
            assert x["capability_adjudication"] == e["capability"] is None
            assert x["capability_verdict_stage"] == e["capability_verdict_stage"]
            assert x["capability_target"] == e["capability_target"]
            assert x["html_counterpart_verdict"] == e["html_counterpart_verdict"]
            assert x["adjudication_reason"] == e["adjudication"]["reason"]
            assert x["html_counterpart_source_refs"] == e["html_counterpart_source_refs"]
            assert x["must_fix_before_wiring"] == e["capability_target_blocked_by"]
            assert x["host_component"] == e["host_path"]

    def test_plan_reason_is_not_circular(self, deletion_plan: dict, paradigm: dict) -> None:
        ap1 = next(
            a for a in paradigm["adjudication_criteria"]["anti_patterns"] if a["id"] == "AP-1")
        markers = ap1["circular_reason_markers"]
        for x in deletion_plan["entries"]:
            if x["capability_adjudication"] != "single_onlyoffice":
                continue
            hit = [m for m in markers if m in x["adjudication_reason"]]
            assert not hit, f"{x['entry_id']}: 循环论证 {hit}"

    def test_plan_two_deletion_classes_partition_the_thirteen_files(
        self, deletion_plan: dict
    ) -> None:
        orphan = {m["file"] for m in deletion_plan["orphan_dual_mode_to_delete"]["modules"]}
        live = {m["file"] for m in deletion_plan["live_dual_mode_to_rewire_then_delete"]["modules"]}
        entry_scoped: set[str] = set()
        for e in deletion_plan["entries"]:
            entry_scoped |= set(e["legacy_composables_to_delete"])
        assert orphan & live == set()
        assert len(orphan) == 7 and len(live) == 6
        assert entry_scoped == live, (
            "entry 范围内的删除对象应恰是 6 个在用 composable（orphan 走独立节）"
        )
        assert len(orphan | live) == 13

    def test_plan_orphans_can_be_deleted_before_step_9(self, deletion_plan: dict) -> None:
        o = deletion_plan["orphan_dual_mode_to_delete"]
        assert o["can_delete_before_step_9"] is True
        assert o["ac_reference"] == "1.7"
        assert deletion_plan["live_dual_mode_to_rewire_then_delete"][
            "can_delete_before_step_9"] is False
        assert o["orphan_barrels"] == []
        assert "os.path.exists" in o["orphan_barrels_note"] or "不存在" in o["orphan_barrels_note"]

    def test_plan_shared_base_numbers_agree_with_the_slice(
        self, deletion_plan: dict, manifest_slice: dict
    ) -> None:
        p = deletion_plan["shared_base_preserved"]
        s = manifest_slice["orphan_dual_mode_inventory"]["shared_base"]
        assert p["file"] == s["path"]
        assert p["statement_position_consumers"] == s["statement_position_consumers"] == 29
        assert p["k_cycle_consumers"] == s["k_cycle_contribution"] == 0
        assert p["remaining_after_this_plan"] == s["remaining_after_k_cycle_work"] == 29
        assert "29 − 0" in p["remaining_after_this_plan_recipe"] or "不变" in p[
            "remaining_after_this_plan_recipe"]

    def test_plan_orphan_modules_mirror_the_slice_inventory(
        self, deletion_plan: dict, manifest_slice: dict
    ) -> None:
        a = deletion_plan["orphan_dual_mode_to_delete"]["modules"]
        b = manifest_slice["orphan_dual_mode_inventory"]["modules"]
        assert [m["file"] for m in a] == [m["file"] for m in b]
        assert [m["lines"] for m in a] == [m["lines"] for m in b]
        assert [m["localStorage_prefix"] for m in a] == [m["localStorage_prefix"] for m in b]

    def test_plan_must_not_wire_to_targets_are_the_orphan_twins(
        self, deletion_plan: dict
    ) -> None:
        """改线时最容易误接的就是 orphan 孪生 —— 必须逐条点名。"""
        orphan = {m["file"] for m in deletion_plan["orphan_dual_mode_to_delete"]["modules"]}
        named: set[str] = set()
        for e in deletion_plan["entries"]:
            for t in e["must_not_wire_to"]:
                named.add(t)
                assert (ROOT / t).exists(), f"must_not_wire_to 指向不存在的文件 {t}"
        assert named == orphan, (
            f"must_not_wire_to 未覆盖全部 orphan 孪生：漏 {sorted(orphan - named)}，"
            f"多 {sorted(named - orphan)}"
        )

    def test_plan_host_inlined_blocks_are_real(self, deletion_plan: dict) -> None:
        n = 0
        for e in deletion_plan["entries"]:
            ref = e["host_inlined_block_to_remove"]
            if ref is None:
                continue
            n += 1
            line = _line_at(ref)
            assert "const dualMode = ((" in line, (
                f"{e['entry_id']}: host_inlined_block_to_remove 那一行不是内联 IIFE"
                f"（实际 {line.strip()!r}）"
            )
        assert n == deletion_plan["counters"]["host_inlined_blocks_to_remove"] == 7

    def test_plan_counters_recompute(self, deletion_plan: dict) -> None:
        c = deletion_plan["counters"]
        E = deletion_plan["entries"]
        assert c["entries"] == len(E) == 13
        assert c["composables_to_delete_entry_scoped"] == sum(
            len(x["legacy_composables_to_delete"]) for x in E) == 6
        assert c["orphan_dual_mode_to_delete"] == len(
            deletion_plan["orphan_dual_mode_to_delete"]["modules"]) == 7
        assert c["live_dual_mode_to_rewire_then_delete"] == len(
            deletion_plan["live_dual_mode_to_rewire_then_delete"]["modules"]) == 6
        assert (c["composables_to_delete_entry_scoped"]
                + c["orphan_dual_mode_to_delete"] == 13)
        assert c["orphan_dual_mode_lines_total"] == sum(
            m["lines"] for m in deletion_plan["orphan_dual_mode_to_delete"]["modules"])
        assert c["orphan_barrel_reexports_to_drop"] == 0
        assert c["hosts_with_inlined_second_implementation"] == sum(
            1 for x in E if x["host_inlined_block_to_remove"]) == 7
        assert (c["hosts_with_inlined_second_implementation"]
                + c["hosts_with_dedicated_composable"] == 13)
        assert c["host_inlined_blocks_to_remove"] == 7
        assert c["localStorage_prefix_declaration_sites"] == 16
        assert c["entries_with_extra_ui_action_required"] == sum(
            1 for x in E if x["extra_ui_action_required"]) == 13
        assert c["entries_with_must_fix_before_wiring"] == sum(
            1 for x in E if x["must_fix_before_wiring"]) == 13
        assert c["entries_with_must_not_wire_to"] == sum(
            1 for x in E if x["must_not_wire_to"]) == 7
        assert c["shared_base_consumers_before"] == c["shared_base_consumers_after"] == 29
        assert c["test_files_requiring_sync_change"] == 0
        assert c["counting_recipe"].strip()

    def test_plan_has_no_pilot_branch(self, deletion_plan: dict) -> None:
        assert deletion_plan["pilot_already_migrated"] is None
        assert deletion_plan["pilot_already_migrated_note"].strip()
        for owner in PILOT_CONTRACT_OWNERS:
            assert owner in deletion_plan["pilot_already_migrated_note"], (
                f"pilot 说明里没有点名 {owner} —— 归属没有实证"
            )

    def test_plan_step5_applicability_differs_from_j(self, deletion_plan: dict) -> None:
        why = deletion_plan["why_this_plan_has_two_deletion_object_classes"]
        assert "why_step_5_applies_here_unlike_j" in why
        assert "16" in why["why_step_5_applies_here_unlike_j"]
        assert why["verification_recipe"].strip()


# ════════════════════════════════════════════════════════════════════════════
def _load_validator():
    """加载范式点名的校验器（复用，不另写一份 —— 两份实现漂移时谁都不红）。"""
    path = pathlib.Path(__file__).with_name("test_migration_paradigm_contract.py")
    spec = importlib.util.spec_from_file_location("_task53_validator_host", path)
    assert spec and spec.loader
    module = importlib.util.module_from_spec(spec)
    spec.loader.exec_module(module)
    return module


class TestParadigmCompliance:
    def test_the_slice_passes_the_paradigm_nominated_validator(self, manifest_slice: dict) -> None:
        v = _load_validator().validate_slice_against_schema(manifest_slice)
        assert v == [], "范式校验器报违规：\n" + "\n".join(f"  - {x}" for x in v)

    def test_the_validator_really_catches_a_missing_pending_field(
        self, manifest_slice: dict
    ) -> None:
        """反向自检：抽掉待裁决三字段之一必须打红，且说明点名该字段。"""
        mod = _load_validator()
        for field in ("capability_verdict_stage", "capability_target",
                      "capability_target_blocked_by"):
            broken = json.loads(json.dumps(manifest_slice))
            broken["independent_entries"][0].pop(field)
            v = mod.validate_slice_against_schema(broken)
            assert any(field in x for x in v), (
                f"抽掉 `{field}` 后校验器没有点名它 —— SR-3 右支形同虚设"
            )

    def test_the_validator_really_catches_a_counter_mismatch(self, manifest_slice: dict) -> None:
        mod = _load_validator()
        broken = json.loads(json.dumps(manifest_slice))
        broken["honest_adjudication_summary"]["slice_counters"]["unadjudicated"] = 0
        v = mod.validate_slice_against_schema(broken)
        assert any("SR-9" in x for x in v), "把待裁决计数改成 0 竟没打红 —— SR-9 是后门"

    def test_the_validator_really_catches_a_forbidden_identity_kind(
        self, manifest_slice: dict
    ) -> None:
        mod = _load_validator()
        broken = json.loads(json.dumps(manifest_slice))
        broken["dynamic_row_identity"]["tables"][0]["row_identity"]["kind"] = "array_index"
        v = mod.validate_slice_against_schema(broken)
        assert any("forbidden_identity_kinds" in x for x in v)

    def test_the_validator_really_catches_a_single_with_identity(
        self, manifest_slice: dict
    ) -> None:
        """SR-5 / AP-5：裁 single 的 entry 挂 adapter 必须打红。"""
        mod = _load_validator()
        broken = json.loads(json.dumps(manifest_slice))
        e = broken["independent_entries"][0]
        e["capability"] = "single_onlyoffice"
        e["adjudication"]["honest_capability"] = "single_onlyoffice"
        e["adapter_id"] = "fake-adapter"
        broken["honest_adjudication_summary"]["adjudicated_as_single_onlyoffice"] = 1
        broken["honest_adjudication_summary"]["slice_counters"]["unadjudicated"] = 12
        v = mod.validate_slice_against_schema(broken)
        assert any("SR-5" in x for x in v), "single_* 挂 adapter 竟没打红"

    def test_paradigm_refs_and_task_number_are_right(self, manifest_slice: dict) -> None:
        assert manifest_slice["task"] == "Task 53"
        assert manifest_slice["schema_version"] == "manifest-slice:v1"
        assert manifest_slice["source_manifest"] == FULL_MANIFEST_PATH.relative_to(
            ROOT).as_posix()
        assert manifest_slice["paradigm_ref"] == PARADIGM_PATH.relative_to(ROOT).as_posix()
        assert manifest_slice["frozen_at"].startswith("Task 53")

    def test_task48_extra_entry_fields_are_present(
        self, manifest_slice: dict, paradigm: dict
    ) -> None:
        extra = paradigm["slice_schema"]["effective_from_task_48"]["required_entry"]
        for e in manifest_slice["independent_entries"]:
            for f in extra:
                assert f in e, f"{e['entry_id']}: 缺 Task 48 起必填的 `{f}`"

    def test_paradigm_bytes_are_untouched_by_this_task(self, paradigm: dict) -> None:
        """本任务不改范式字节：`paradigm` 块的 digest 双向锁死。"""
        reg = paradigm["paradigm_registry"]["paradigms"][0]
        assert reg["alias"] == "legacy_deletion_paradigm"
        recomputed = hashlib.sha256(
            json.dumps(paradigm["paradigm"], ensure_ascii=False, indent=2,
                       sort_keys=True).encode("utf-8")
        ).hexdigest()
        assert recomputed == reg["frozen_canonical_sha256"], (
            "🔴 `paradigm` 块的字节变了 —— 要么本任务误改了范式，要么范式 owner 改了但没更 digest"
        )
        assert len(paradigm["paradigm"]["steps"]) == reg["step_count"] == 7
        assert paradigm["adjudication_criteria"]["capability_enum"] == list(CAPABILITY_ENUM), (
            "capability_enum 被扩了 —— AC 1.3 原文派生，不得改"
        )

    def test_new_sections_are_declared_as_non_conflicting(self, manifest_slice: dict) -> None:
        conflict = manifest_slice["paradigm_schema_conflict"]
        assert conflict["residual_inconsistency"] is None
        assert conflict["known_red_it_causes"]["status"] == "none"
        ns = conflict["new_section_not_in_schema"]
        for section in ns["sections"]:
            assert section in manifest_slice, f"声明的新增节 {section} 实际不在 slice 里"
        assert conflict["paradigm_bytes_untouched"].strip()
        assert conflict["sibling_slice_conflicts"].strip()
        locked_by = conflict["known_red_it_causes"]["locked_by"]
        assert "test_task53_k_cycle_migration.py" in locked_by

    def test_unjudged_slices_registry_stays_empty(self) -> None:
        """本 slice 不得给 `_UNJUDGED_SLICES` 添一条 —— 它必须直接过校验。"""
        path = pathlib.Path(__file__).with_name("test_slice_schema_validator_coverage.py")
        src = path.read_text(encoding="utf-8")
        m = re.search(r"_UNJUDGED_SLICES:\s*dict\[str,\s*str\]\s*=\s*(\{[^}]*\})", src)
        assert m, "找不到 _UNJUDGED_SLICES 声明"
        assert m.group(1).strip() == "{}", (
            f"_UNJUDGED_SLICES 不再为空：{m.group(1)} —— 本任务不得靠登记豁免过关"
        )

    def test_requirements_and_properties_are_declared(self, manifest_slice: dict) -> None:
        reqs = set(manifest_slice["requirements_covered"])
        assert {"6.1", "6.2", "6.6", "12.1", "12.4", "12.10", "12.11", "12.12", "14.1"} <= reqs
        props = " ".join(manifest_slice["properties_verified"])
        for n in (20, 24, 69, 70):
            assert f"Property {n}" in props, f"tasks.md 点名的 Property {n} 未声明"
        pd = manifest_slice["property_denominators"]
        for key in ("property_20", "property_24", "property_69", "property_70"):
            assert key in pd, f"{key} 缺分母声明"
            assert pd[key]["k_cycle_denominator"].strip()
            assert "not_claimed_passing" in pd[key]

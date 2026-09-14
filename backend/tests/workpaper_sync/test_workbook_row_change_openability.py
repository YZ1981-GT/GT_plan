# -*- coding: utf-8 -*-
"""Task 25 —— 传播产物的真实可打开性与求值层级。

spec: excel-workbook-wide-row-change-propagation / Wave 5 Task 25
Requirements: 8.4
Properties: **P31**

═══ 三层判据，强度递增 ═══════════════════════════════════════════════════

1. **本仓库的安全/容量门** —— `validate_ooxml_artifact`（`errors` 必须为空）。
   只证明「没触发我们自己的门」。
2. **openpyxl 能加载** —— 一个独立的第三方 OOXML 实现认得它，且传播后的公式文本
   在它眼里就是传播后的样子。
3. 🔴 **OnlyOffice 9.4 的真实文档引擎 `x2t`** —— xlsx → bin → xlsx 双向走通，
   且回转产物里被传播的引用**逐字保留**。这一层才是「真实引擎认得」。

═══ 🔴 AC 8.4 的「公式求值正确」被诚实拆成两件事 ═══════════════════════

AC 8.4 原文是「可打开**且公式求值正确**」。C 在 2026-09-05 实测了 x2t 的求值行为，
结论是**它只搬运 cached `<v>`，不重算**：

    往受管表的目标格 `'明细表D2-2'!AC28` 写入可辨认值 424242，
    再送 x2t 走 xlsx→bin→xlsx：
      · 目标格本体      → <c r="AC28"><v>424242</v></c>   （值被保留）
      · 引用它的两个格  → <v>0</v>                        （**没有**变成 424242）
    两次回转（写标记前 / 写标记后）引用格的 `<v>` 完全相同 ⇒ 未重算。

所以本文件把 AC 8.4 拆成：

| 层 | 能验什么 | 本文件的处置 |
|---|---|---|
| 结构 | 引擎解析→序列化→写回后，传播后的引用逐字保留 | ✅ 断言（可复算） |
| 求值 | Excel 重算后指向的是新行的数据 | ⚠ **UNVERIFIABLE**，需真 Excel/人工 |

🔴 **不用 x2t 的「转换成功」冒充「求值正确」。** AC 8.4 明令「不得用 fixture 冒充」——
这次的冒充者不是 fixture，而是一个**看起来更权威**的真实引擎，因此更容易蒙混过关：
x2t 返回 rc=0 且产物里公式文本正确，很容易被写成「真实引擎验证求值通过」。
它证明的是**引用被正确保留**，不是**引用求出的值正确**。两者差一个重算引擎。

环境不可得时 **skip 并标 UNVERIFIABLE**（AC 8.4），并在 skip 理由里写明缺什么。
"""

from __future__ import annotations

import io
import os
import re
import shutil
import subprocess
import sys
import zipfile
from pathlib import Path
from typing import Any, Final

import pytest

_REPO = Path(__file__).resolve().parents[3]
_BACKEND = _REPO / "backend"
if str(_BACKEND) not in sys.path:  # pragma: no cover - import 环境自举
    sys.path.insert(0, str(_BACKEND))
os.environ.setdefault("DB_DISABLE_SSL", "True")

from app.services.excel_structure_fingerprint import (  # noqa: E402
    _normalise_part,
    _parse_workbook_xml,
)
from app.services.workpaper_sync import excel_workbook_row_change as N1  # noqa: E402
from app.services.workpaper_sync.ooxml_security import (  # noqa: E402
    validate_ooxml_artifact,
)

TEMPLATE_ROOT: Final[Path] = _BACKEND / "wp_templates"

#: 🔴 首要判据载体 = D2（Wave 0 Gate 1 裁决：今天唯一既有已审契约又有真实跨 sheet 引用）。
#:    文件名含**两个连续空格**，别按 spec 正文的省略写法拼路径。
D2_REL: Final[str] = "D/D2-1至D2-4  应收账款- 审定表明细表（Leap-常规程序）.xlsx"
D2_SHEET: Final[str] = "明细表D2-2"
D2_AT: Final[int] = 26
D2_COUNT: Final[int] = 2
D2_STYLE_FROM: Final[int] = 25
D2_REGION: Final[tuple[int, int]] = (13, 25)

#: OnlyOffice 容器名与引擎路径 —— 现探，不写死假设。
_OO_CONTAINER: Final[str] = os.environ.get("GT_ONLYOFFICE_CONTAINER", "audit-onlyoffice")
_X2T: Final[str] = "/var/www/onlyoffice/documentserver/server/FileConverter/bin/x2t"

#: x2t 的格式号：8192 = 内部 bin，257 = xlsx。
_FMT_BIN: Final[int] = 8192
_FMT_XLSX: Final[int] = 257


# ═══════════════════════════════════════════════════════════════════════════
# 1. 环境探测 —— 取不到就 UNVERIFIABLE，不冒充
# ═══════════════════════════════════════════════════════════════════════════


def _docker() -> str | None:
    return shutil.which("docker")


def _onlyoffice_engine() -> tuple[bool, str]:
    """`(可用, 版本号或 UNVERIFIABLE 理由)` —— 现探容器与二进制。

    与 `test_excel_row_insertion_openability._onlyoffice_engine` 同款约定（B 的范式），
    刻意不 import 它：那份是上游 spec 的测试文件，跨 spec import 会让两个 spec 的
    测试互相成为对方的前置依赖。
    """
    docker = _docker()
    if docker is None:
        return False, "UNVERIFIABLE: 宿主上没有 docker CLI"
    probe = subprocess.run(  # noqa: S603 - 固定参数，无 shell
        [docker, "exec", _OO_CONTAINER, "sh", "-c", f"test -x {_X2T} && {_X2T} 2>&1 | head -4"],
        capture_output=True,
        text=True,
        timeout=120,
    )
    if probe.returncode != 0:
        stderr = (probe.stderr or "").strip()
        # 🔴 把「守护进程挂了」与「容器里没这个文件」分开报。
        #    2026-09-05 实测踩到过：Docker Desktop 中途挂掉时 stderr 是
        #    `Docker Desktop is unable to start`，而原来的理由一律写成「取不到 x2t」
        #    ⇒ 下一个人会去容器里找文件，而真正该做的是把 Docker 拉起来。
        #    skip 理由不精确 = 把人导向错误的排查方向。
        if "unable to start" in stderr or "daemon" in stderr.lower():
            return False, (
                f"UNVERIFIABLE: Docker 守护进程不可用（{stderr[:120]}）—— "
                "不是容器或 x2t 的问题，先把 Docker 拉起来再复算本层"
            )
        return False, (
            f"UNVERIFIABLE: 容器 {_OO_CONTAINER} 里取不到 {_X2T}"
            f"（rc={probe.returncode}, err={stderr[:120]}）"
        )
    version = re.search(r"Version:\s*([\d.]+)", probe.stdout or "")
    if version is None:
        return False, f"UNVERIFIABLE: x2t 没报版本号（stdout={(probe.stdout or '')[:120]!r}）"
    return True, version.group(1)


def _x2t_roundtrip(payload: bytes, *, label: str) -> tuple[int, bytes | None]:
    """xlsx → bin → xlsx，返回 `(最后一步退出码, 回转产物字节)`。

    🔴 必须走 **params.xml** 协议：直接 `x2t in.xlsx out.bin` 会报
    `Couldn't create temp folder`（实测），那不是权限问题而是调用协议不对 ——
    误判成权限问题会让人去 chown，然后依然失败。

    子进程一律 `subprocess.run([...])` 传数组，路径含中文与空格。
    """
    docker = _docker()
    assert docker is not None, "调用方保证 docker 可得"
    remote = f"/tmp/gt_wrc_open_{label}"
    subprocess.run(  # noqa: S603
        [docker, "exec", _OO_CONTAINER, "sh", "-c", f"rm -rf {remote} && mkdir -p {remote}"],
        check=True,
        capture_output=True,
        timeout=60,
    )
    local = Path(os.environ.get("TEMP", "/tmp")) / f"gt_wrc_{label}.xlsx"
    try:
        local.write_bytes(payload)
        subprocess.run(  # noqa: S603
            [docker, "cp", str(local), f"{_OO_CONTAINER}:{remote}/in.xlsx"],
            check=True,
            capture_output=True,
            timeout=180,
        )
        rc = -1
        for frm, to, fmt in (("in.xlsx", "out.bin", _FMT_BIN), ("out.bin", "back.xlsx", _FMT_XLSX)):
            params = (
                '<?xml version="1.0" encoding="utf-8"?>'
                "<TaskQueueDataConvert>"
                f"<m_sFileFrom>{remote}/{frm}</m_sFileFrom>"
                f"<m_sFileTo>{remote}/{to}</m_sFileTo>"
                f"<m_nFormatTo>{fmt}</m_nFormatTo>"
                "</TaskQueueDataConvert>"
            )
            subprocess.run(  # noqa: S603
                [
                    docker, "exec", _OO_CONTAINER, "sh", "-c",
                    f"cat > {remote}/p.xml <<'XMLEOF'\n{params}\nXMLEOF",
                ],
                check=True,
                capture_output=True,
                timeout=60,
            )
            proc = subprocess.run(  # noqa: S603
                [
                    docker, "exec", _OO_CONTAINER, "sh", "-c",
                    f"cd {remote} && {_X2T} p.xml >/dev/null 2>&1; echo rc=$?; "
                    f"test -s {remote}/{to} && echo size=$(stat -c %s {remote}/{to}) || echo size=0",
                ],
                capture_output=True,
                text=True,
                timeout=300,
            )
            blob = (proc.stdout or "") + (proc.stderr or "")
            hit = re.search(r"rc=(-?\d+)", blob)
            size = re.search(r"size=(\d+)", blob)
            assert hit is not None and size is not None, f"取不到 rc/size: {blob[:300]!r}"
            rc = int(hit.group(1))
            if rc != 0 or int(size.group(1)) == 0:
                return rc, None
        back = Path(os.environ.get("TEMP", "/tmp")) / f"gt_wrc_{label}_back.xlsx"
        subprocess.run(  # noqa: S603
            [docker, "cp", f"{_OO_CONTAINER}:{remote}/back.xlsx", str(back)],
            check=True,
            capture_output=True,
            timeout=180,
        )
        data = back.read_bytes() if back.is_file() else None
        back.unlink(missing_ok=True)
        return rc, data
    finally:
        local.unlink(missing_ok=True)
        subprocess.run(  # noqa: S603
            [docker, "exec", _OO_CONTAINER, "sh", "-c", f"rm -rf {remote}"],
            capture_output=True,
            timeout=60,
        )


# ═══════════════════════════════════════════════════════════════════════════
# 2. 真实传播产物 —— 走生产路径，不手搓
# ═══════════════════════════════════════════════════════════════════════════


def _rewrite_zip_xml(data: bytes, old: str, new: str) -> bytes:
    """在 zip **内部各 XML 部件**里做文本替换，返回重打包后的字节。

    🔴 存在的理由是变异检验：xlsx 里的 XML 是 DEFLATE 压缩的，对 zip 容器做
    `bytes.replace` 根本碰不到内容 —— 实测过，那样的变异用例必判 GREEN 并被
    误读成「守卫抓不住引用错行」。注入错误产物必须解包→改 XML→重打包。
    """
    with zipfile.ZipFile(io.BytesIO(data)) as zf:
        entries = {n: zf.read(n) for n in zf.namelist()}
    for name, blob in list(entries.items()):
        if not name.endswith(".xml"):
            continue
        text = blob.decode("utf-8", "replace")
        if old in text:
            entries[name] = text.replace(old, new).encode("utf-8")
    buf = io.BytesIO()
    with zipfile.ZipFile(buf, "w", zipfile.ZIP_DEFLATED) as out:
        for name, blob in entries.items():
            out.writestr(name, blob)
    return buf.getvalue()


def _sheet_parts(data: bytes) -> tuple[dict[str, str], list[dict[str, Any]]]:
    with zipfile.ZipFile(io.BytesIO(data)) as zf:
        sheets, defined = _parse_workbook_xml(zf)
        return {s["name"]: _normalise_part(s["rel_target"]) for s in sheets}, list(defined)


@pytest.fixture(scope="module")
def d2_propagated() -> tuple[bytes, bytes, N1.WorkbookRowChangePlan, N1.PropagationReport]:
    """`(原字节, 传播后字节, 计划, 报告)` —— 全程走生产 API。"""
    path = TEMPLATE_ROOT / D2_REL
    if not path.is_file():  # pragma: no cover - 模板缺失时不冒充通过
        pytest.skip(f"UNVERIFIABLE: D2 权威模板不在磁盘上：{D2_REL}")
    data = path.read_bytes()
    parts, defined = _sheet_parts(data)
    with zipfile.ZipFile(io.BytesIO(data)) as zf:
        scan = N1.scan_reference_carriers(
            zf, target_sheet=D2_SHEET, sheet_parts=parts, defined_names=defined
        )
    plan = N1.build_insert_plan(
        scan,
        managed_sheet_name=D2_SHEET,
        managed_sheet_part=parts[D2_SHEET],
        at=D2_AT,
        count=D2_COUNT,
        style_from=D2_STYLE_FROM,
        region_first_row=D2_REGION[0],
        region_last_row=D2_REGION[1],
    )
    produced, report = N1.apply_workbook_row_change(data, plan, sheet_parts=parts)
    return data, produced, plan, report


@pytest.fixture(scope="module")
def moved_entries(
    d2_propagated: tuple[bytes, bytes, N1.WorkbookRowChangePlan, N1.PropagationReport],
) -> tuple[N1.PropagationEntry, ...]:
    """行号**真的变了**的传播条目。

    🔴 分母判据的落点：若这个集合是空的，后面所有「引用被正确保留」的断言都会在空集上
    恒真。实测 D2 上 26/26 条全部行号有变（`$13:$25` 那类在插入点之上的区间根本没进
    传播清单 —— 扫描器判定精准，不是「进了清单但没改」）。
    """
    _data, _produced, plan, _report = d2_propagated
    moved = tuple(e for e in plan.propagations if e.row_before != e.row_after)
    assert moved, "没有任何条目行号发生变化 —— 判据会在空集上恒真"
    return moved


class TestLayer1OurOwnGates:
    """第一层：本仓库的安全/容量门。只证明「没触发我们自己的门」。

    **Validates: Requirements 8.4**
    """

    def test_validate_ooxml_artifact_reports_no_errors(
        self,
        d2_propagated: tuple[bytes, bytes, N1.WorkbookRowChangePlan, N1.PropagationReport],
        tmp_path: Path,
    ) -> None:
        _data, produced, _plan, _report = d2_propagated
        target = tmp_path / "propagated.xlsx"
        target.write_bytes(produced)
        report = validate_ooxml_artifact(target, document_type="xlsx")
        assert report is not None
        payload = report.as_dict() if hasattr(report, "as_dict") else {}
        errors = payload.get("errors") or []
        assert not errors, f"传播产物触发了本仓库的 OOXML 门: {errors}"

    def test_propagation_actually_happened(
        self,
        d2_propagated: tuple[bytes, bytes, N1.WorkbookRowChangePlan, N1.PropagationReport],
        moved_entries: tuple[N1.PropagationEntry, ...],
    ) -> None:
        """🔴 分母判据：产物**真的**被传播过，而不是「原样复制也能通过前两层」。

        缺这一条时，把 `apply_workbook_row_change` 换成 `lambda data, *a, **k: (data, report)`
        依然全绿 —— 那是本 spec 反复登记的假绿形态（「什么都没做也通过」）。
        """
        data, produced, _plan, report = d2_propagated
        assert produced != data, "产物与原字节相同 —— 传播根本没发生"
        # 🔴 计数器名从生产的 `CARRIER_COUNTER_NAMES` 取，不在判据里写死字段名 ——
        #    写死会让「载体清单加一项」这件事在本判据上静默通过。
        counter = N1.CARRIER_COUNTER_NAMES["formula"]
        assert getattr(report, counter) > 0, f"公式传播计数为 0: {report.as_dict()}"
        assert len(moved_entries) >= 20, (
            f"行号有变的条目只有 {len(moved_entries)} 条，D2 实测应有 26 条 —— "
            "分母偏少说明扫描面被缩小了"
        )


class TestLayer2ThirdPartyImplementation:
    """第二层：openpyxl（独立的第三方 OOXML 实现）能加载，且看到的是传播**后**的公式。

    **Validates: Requirements 8.4**
    """

    def test_openpyxl_loads_the_propagated_artifact(
        self,
        d2_propagated: tuple[bytes, bytes, N1.WorkbookRowChangePlan, N1.PropagationReport],
    ) -> None:
        _data, produced, _plan, _report = d2_propagated
        from openpyxl import load_workbook

        wb = load_workbook(io.BytesIO(produced))
        try:
            assert D2_SHEET in wb.sheetnames, (
                f"受管 sheet {D2_SHEET!r} 在 openpyxl 眼里不存在，实有 {wb.sheetnames}"
            )
        finally:
            wb.close()

    def test_openpyxl_sees_shifted_references_not_stale_ones(
        self,
        d2_propagated: tuple[bytes, bytes, N1.WorkbookRowChangePlan, N1.PropagationReport],
        moved_entries: tuple[N1.PropagationEntry, ...],
    ) -> None:
        """🔴 独立实现读出来的公式文本 == 声明的 `ref_after`，且**不含** `ref_before`。

        两个方向都断言：只断言「含 after」时，一个把公式写成 `after + before` 的实现
        也能通过；只断言「不含 before」时，把公式整条删掉也能通过。
        """
        _data, produced, _plan, _report = d2_propagated
        from openpyxl import load_workbook

        wb = load_workbook(io.BytesIO(produced))
        try:
            checked = 0
            for entry in moved_entries:
                if entry.carrier != "formula":
                    continue
                coord = entry.locator.split("#")[0]
                sheet_name = _sheet_name_of_part(produced, entry.part)
                if sheet_name is None or sheet_name not in wb.sheetnames:
                    continue
                value = wb[sheet_name][coord].value
                if not isinstance(value, str):
                    continue
                assert entry.ref_after in value, (
                    f"{sheet_name}!{coord} 的公式里没有传播后的引用 {entry.ref_after!r}：{value!r}"
                )
                # 🔴 比**坐标片段**而不是完整 `ref_before`。
                #
                #    变异实测（M3）：注入 `!AC28+0*AC26` 让传播前后的坐标**并存**时，
                #    用完整 `ref_before`（`'明细表D2-2'!AC26`）做子串检查**抓不住** ——
                #    注入的是裸 `AC26`，没有 sheet 前缀。而「同一 sheet 的裸引用回退」
                #    恰恰是最可能的错法（改写器少加了一次前缀就退化成这样）。
                #    改比坐标片段后该变异立刻打红，且与第三层的口径一致。
                before_coord = _coord_of(entry.ref_before)
                after_coord = _coord_of(entry.ref_after)
                if before_coord and after_coord and before_coord != after_coord:
                    assert not re.search(rf"(?<![A-Z0-9]){re.escape(before_coord)}(?![0-9])", value), (
                        f"{sheet_name}!{coord} 的公式里还留着传播**前**的坐标 "
                        f"{before_coord!r}（after={after_coord!r}）：{value!r}"
                    )
                checked += 1
            assert checked >= 10, f"只核到 {checked} 处公式，分母偏少（防空集恒真）"
        finally:
            wb.close()


def _sheet_name_of_part(data: bytes, part: str) -> str | None:
    with zipfile.ZipFile(io.BytesIO(data)) as zf:
        sheets, _ = _parse_workbook_xml(zf)
    for sheet in sheets:
        if _normalise_part(sheet["rel_target"]) == part:
            return str(sheet["name"])
    return None


# ═══════════════════════════════════════════════════════════════════════════
# 3. 第三层：OnlyOffice 9.4 的真实文档引擎
# ═══════════════════════════════════════════════════════════════════════════


def _cells_referencing(data: bytes, *, sheet_name: str, coord: str) -> list[tuple[str, str]]:
    """`[(part, 单元格 XML)]` —— 公式文本里同时含 `sheet_name` 与 `coord` 的格。

    🔴 按**公式内容**定位，不按坐标猜：x2t 回转会重排 sheet 顺序、也会把空格压掉，
    原来的 `审定表D2-1!B8` 在回转产物里可能变成别的 part 的 `B7`。按坐标找会找不到，
    然后被写成「引用消失了」的假红。
    """
    found: list[tuple[str, str]] = []
    needle = sheet_name
    with zipfile.ZipFile(io.BytesIO(data)) as zf:
        for name in zf.namelist():
            if not name.startswith("xl/worksheets/sheet"):
                continue
            xml = zf.read(name).decode("utf-8", "replace")
            for match in re.finditer(r"<c r=\"[A-Z]+\d+\"[^>]*>.*?</c>", xml, re.S):
                blob = match.group(0)
                if needle in blob and coord in blob:
                    found.append((name, blob))
    return found


class TestLayer3RealOnlyOfficeEngine:
    """第三层：OnlyOffice 9.4 的 `x2t` 双向走通，且传播后的引用逐字保留。

    **Validates: Requirements 8.4**

    环境不可得时 **skip 并标 UNVERIFIABLE**，绝不用 fixture 冒充。
    """

    def test_engine_is_94_or_skip_as_unverifiable(self) -> None:
        """引擎必须是 OnlyOffice **9.4**。

        🔴 判据是**按点分段比对**，不是 `startswith("9.4")`。变异实测（I6）：注入
        `return True, "9.4-FAKE"` 时前缀匹配照样通过 —— 而真实风险是引擎升到
        `9.40`（`"9.40".startswith("9.4")` 为真）会被当成 9.4 静默放过，
        届时 identity 载体与转换行为的既有实测结论全部失效却无人知晓。
        """
        available, detail = _onlyoffice_engine()
        if not available:
            pytest.skip(detail)
        parts = detail.split(".")
        assert len(parts) >= 2 and parts[0].isdigit() and parts[1].isdigit(), (
            f"引擎版本号形态非法（应形如 `9.4.0.129`）：{detail!r} —— "
            "拿不到可比对的版本号时不得当成通过"
        )
        assert (int(parts[0]), int(parts[1])) == (9, 4), (
            f"本平台的目标引擎是 OnlyOffice 9.4，实测 {detail} —— "
            "换了版本，identity 载体与转换行为的既有实测结论都需重新取证"
        )

    def test_x2t_roundtrips_the_propagated_artifact(
        self,
        d2_propagated: tuple[bytes, bytes, N1.WorkbookRowChangePlan, N1.PropagationReport],
    ) -> None:
        """🔴 带对照组：**未传播**的同一份 substrate 自己先要能走通。

        缺对照时，若容器坏了（任何输入都 rc≠0），本条会把「环境坏了」误报成
        「传播产物坏了」。
        """
        available, detail = _onlyoffice_engine()
        if not available:
            pytest.skip(detail)
        data, produced, _plan, _report = d2_propagated

        control_rc, control_back = _x2t_roundtrip(data, label="control")
        assert control_rc == 0 and control_back, (
            f"对照（未传播的原模板）自己就走不通 rc={control_rc} —— 环境问题，非产物问题"
        )

        rc, back = _x2t_roundtrip(produced, label="propagated")
        assert rc == 0, f"OnlyOffice 9.4 的 x2t 打不开传播产物（rc={rc}）"
        assert back, "x2t 报成功却没产出回转字节"
        assert len(back) >= len(control_back) * 0.8, (
            f"回转产物 {len(back)} 字节远小于对照 {len(control_back)} —— 疑似丢内容"
        )

    def test_real_engine_preserves_shifted_references_verbatim(
        self,
        d2_propagated: tuple[bytes, bytes, N1.WorkbookRowChangePlan, N1.PropagationReport],
        moved_entries: tuple[N1.PropagationEntry, ...],
    ) -> None:
        """🔴 本文件最强的一条：真实引擎解析→序列化→写回后，**传播后**的行号仍在。

        判据形态是「传播后的坐标在、传播前的坐标不在」两个方向都断言 —— 只断言前者时，
        一个把两份引用都写进去的实现也能过。

        实测样本（2026-09-05）：`'明细表D2-2'!AC26` 传播成 `AC28`，回转产物里是
        `<f>&apos;明细表D2-2&apos;!AC28</f>` —— 注意引擎把单引号写成了 `&apos;` 实体，
        所以判据只能比**坐标片段**（`AC28`），不能整条 `ref_after` 做子串匹配。
        """
        available, detail = _onlyoffice_engine()
        if not available:
            pytest.skip(detail)
        _data, produced, _plan, _report = d2_propagated
        rc, back = _x2t_roundtrip(produced, label="refcheck")
        assert rc == 0 and back, f"x2t 回转失败 rc={rc}"

        checked = 0
        for entry in moved_entries:
            if entry.carrier != "formula":
                continue
            after_coord = _coord_of(entry.ref_after)
            before_coord = _coord_of(entry.ref_before)
            if after_coord is None or before_coord is None or after_coord == before_coord:
                continue
            hits = _cells_referencing(back, sheet_name=D2_SHEET, coord=after_coord)
            assert hits, (
                f"回转产物里找不到引用 {D2_SHEET}!{after_coord} 的格 —— "
                f"传播后的引用没被引擎保留（条目 {entry.locator} "
                f"{entry.ref_before} -> {entry.ref_after}）"
            )
            for part, blob in hits:
                assert before_coord not in blob, (
                    f"{part} 的格里同时留着传播**前**的坐标 {before_coord}：{blob[:200]}"
                )
            checked += 1
        assert checked >= 10, f"只核到 {checked} 处，分母偏少（防空集恒真）"


def _coord_of(ref: str) -> str | None:
    """从 `'明细表D2-2'!AC26` 取出 `AC26`；取不到返回 `None`。"""
    if "!" not in ref:
        return None
    tail = ref.rsplit("!", 1)[1]
    return tail or None


# ═══════════════════════════════════════════════════════════════════════════
# 4. 🔴 AC 8.4 的「公式求值正确」—— 诚实拆层，不用真实引擎冒充
# ═══════════════════════════════════════════════════════════════════════════

#: 2026-09-05 C 的实测结论：x2t **不重算**公式。
#:
#: 实验：往受管表目标格 `'明细表D2-2'!AC28` 写入 424242，再送 x2t 走 xlsx→bin→xlsx。
#:   · 目标格本体      → `<c r="AC28"><v>424242</v></c>`（值被搬运保留）
#:   · 引用它的两个格  → `<v>0</v>`（**没有**变成 424242）
#: 写标记前/后两次回转，引用格的 `<v>` 完全相同 ⇒ 未重算。
X2T_RECALCULATES_FORMULAS: Final[bool] = False


class TestFormulaEvaluationIsHonestlyScoped:
    """AC 8.4 的「求值正确」必须与「结构可打开」可区分，且不得用 x2t 冒充。

    **Validates: Requirements 8.4**

    ═══ 为什么这一节必须存在 ═══

    AC 8.4 写的是「可打开**且公式求值正确**」，并明令「不得用 fixture 冒充」。
    第三层拿到的是真实引擎 rc=0 且公式文本正确 —— 这非常容易被写成
    「已用真实 OnlyOffice 验证求值通过」。

    🔴 那是冒充，而且比 fixture 冒充更危险：fixture 一眼假，**真实引擎的成功**
    看起来是最强的证据。但 x2t 是**转换器**不是**计算引擎**：它搬运 cached `<v>`，
    从不重算。所以它能证明「引用被正确保留」，证明不了「引用求出的值正确」。

    两者差一个重算引擎。少了这层区分，一个把所有公式都改成指向空行的实现
    照样能让第三层全绿。
    """

    def test_x2t_is_documented_as_not_recalculating(self) -> None:
        """把「x2t 不重算」这个能力边界钉成判据，而不是写在注释里。

        若哪天 x2t 开始重算了（或换了会重算的引擎），本条会红 —— 那是**好事**：
        届时求值层就能从 UNVERIFIABLE 升级成可验，而升级必须是**有意**的决定。
        """
        assert X2T_RECALCULATES_FORMULAS is False, (
            "X2T_RECALCULATES_FORMULAS 被改成 True —— 若这是实测结论的更新，"
            "请把 test_formula_evaluation_is_unverifiable_by_x2t 改成真正的求值判据；"
            "留着一个说『不可验』的登记同时声明引擎会重算，是自相矛盾的"
        )

    def test_formula_evaluation_is_unverifiable_by_x2t(
        self,
        d2_propagated: tuple[bytes, bytes, N1.WorkbookRowChangePlan, N1.PropagationReport],
        moved_entries: tuple[N1.PropagationEntry, ...],
    ) -> None:
        """🔴 用实验**重新证明**「x2t 不重算」，不是只信上面那个常量。

        做法：造一份变体，把受管表的传播目标格填上可辨认值，两份都送 x2t 回转，
        比对引用格的 cached `<v>`。若两者相同 ⇒ 未重算 ⇒ 求值层登记 UNVERIFIABLE。

        这条是「常量登记」的反向自检：常量可能被人改错，实验不会。
        """
        available, detail = _onlyoffice_engine()
        if not available:
            pytest.skip(f"{detail}｜求值层本就登记为 UNVERIFIABLE，本条只是复算其依据")
        _data, produced, _plan, _report = d2_propagated

        target = next(
            (
                _coord_of(e.ref_after)
                for e in moved_entries
                if e.carrier == "formula" and _coord_of(e.ref_after)
            ),
            None,
        )
        assert target, "找不到任何传播后的目标坐标 —— 上游 fixture 有问题"

        marked = _with_cell_value(produced, sheet_name=D2_SHEET, coord=target, value="424242")
        rc_a, back_a = _x2t_roundtrip(produced, label="eval_plain")
        rc_b, back_b = _x2t_roundtrip(marked, label="eval_marked")
        assert rc_a == 0 and back_a and rc_b == 0 and back_b, (
            f"回转失败 rc_a={rc_a} rc_b={rc_b}"
        )

        vals_a = _cached_values(back_a, sheet_name=D2_SHEET, coord=target)
        vals_b = _cached_values(back_b, sheet_name=D2_SHEET, coord=target)
        assert vals_a, f"找不到引用 {D2_SHEET}!{target} 的格 —— 无从判断是否重算"

        if "424242" in vals_b:
            pytest.fail(
                "x2t 竟然重算了公式（引用格的 cached 值随目标格改变）—— "
                "这是能力边界变化：请把 X2T_RECALCULATES_FORMULAS 改成 True，"
                "并把求值层从 UNVERIFIABLE 升级成真正的求值判据"
            )
        assert vals_a == vals_b, (
            f"引用格的 cached 值在两次回转间变了但不是标记值：{vals_a} vs {vals_b} —— "
            "形态未知，需人工判断"
        )

    def test_unverifiable_layer_is_distinguishable_from_verified(self) -> None:
        """AC 8.4：「环境不可得/能力不足」必须与「验证通过」可区分。

        判据落在**行为**上：把容器名改成不存在的，探测必须返回 `UNVERIFIABLE:` 前缀，
        而不是静默返回可用。
        """
        available, detail = _onlyoffice_engine()
        if _docker() is None:
            pytest.skip("宿主没有 docker，本条无从区分")
        if available:
            assert not detail.startswith("UNVERIFIABLE"), (
                "探测报可用却带 UNVERIFIABLE 前缀 —— 两种状态混在了一起"
            )
        else:
            assert detail.startswith("UNVERIFIABLE:"), (
                f"环境不可得时的理由必须以 `UNVERIFIABLE:` 开头（便于 grep 汇总），实得 {detail!r}"
            )


def _with_cell_value(data: bytes, *, sheet_name: str, coord: str, value: str) -> bytes:
    """返回一份把 `sheet_name!coord` 的内容换成裸值 `value` 的副本（只在内存里）。"""
    parts, _defined = _sheet_parts(data)
    part = parts[sheet_name]
    with zipfile.ZipFile(io.BytesIO(data)) as zf:
        entries = {n: zf.read(n) for n in zf.namelist()}
    xml = entries[part].decode("utf-8", "replace")
    existing = re.search(rf'<c r="{coord}"[^>]*>.*?</c>|<c r="{coord}"[^>]*/>', xml, re.S)
    replacement = f'<c r="{coord}"><v>{value}</v></c>'
    if existing:
        xml = xml.replace(existing.group(0), replacement, 1)
    else:
        row = re.search(rf'<row r="{re.escape(coord.lstrip("ABCDEFGHIJKLMNOPQRSTUVWXYZ"))}"[^>]*>', xml)
        assert row, f"{sheet_name} 上既没有 {coord} 也没有对应的 <row>"
        xml = xml.replace(row.group(0), row.group(0) + replacement, 1)
    entries[part] = xml.encode("utf-8")
    buf = io.BytesIO()
    with zipfile.ZipFile(buf, "w", zipfile.ZIP_DEFLATED) as out:
        for name, blob in entries.items():
            out.writestr(name, blob)
    return buf.getvalue()


def _cached_values(data: bytes, *, sheet_name: str, coord: str) -> list[str]:
    """引用 `sheet_name!coord` 的那些格里的 cached `<v>` 值。"""
    return [
        v
        for _part, blob in _cells_referencing(data, sheet_name=sheet_name, coord=coord)
        for v in re.findall(r"<v>([^<]*)</v>", blob)
    ]

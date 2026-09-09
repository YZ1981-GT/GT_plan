"""G1-1 behavior and in-memory mutation counterexamples; no business DB."""
import ast
import copy
import hashlib
from pathlib import Path
from types import SimpleNamespace as NS
from unittest.mock import AsyncMock

import pytest
from app.services.workpaper_sync import content_mutation as cm
from app.services.workpaper_sync import publish_time_structure_hash as hashes
from app.services.workpaper_sync import contracts
from app.services.workpaper_sync.published_identity_observer import PublishedIdentityObserver
from backend.tests.workpaper_sync.g1_structure_fixture import workbook_fixture
from backend.tests.workpaper_sync.test_task15_content_mutation import xlsx_payload

ROOT = Path(__file__).resolve().parents[3]


def test_final_artifact_and_missing_identity(tmp_path):
    contract = contracts.parse_contract(xlsx_payload())
    data, anchors = workbook_fixture()
    output = tmp_path / "final.xlsx"
    output.write_bytes(data)
    plan = NS(document_type="xlsx", contract=contract, structure_anchors=anchors)
    service = object.__new__(cm.ContentMutationService)
    legacy = NS(structure_hash=hashlib.sha256(data).hexdigest())
    expected = hashes.compute_structure_hash_from_artifact(data=data, contract=contract, anchors=anchors)
    assert service._projection_structure_hash(plan=plan, output=output, materialized=legacy) == expected
    assert expected != legacy.structure_hash
    changed, _ = workbook_fixture(rows=3)
    # Dynamic row count is not itself a declared-field coordinate change.
    assert hashes.compute_structure_hash_from_artifact(data=changed, contract=contract, anchors=anchors) == expected
    from openpyxl import load_workbook
    import io

    wb = load_workbook(io.BytesIO(data))
    del wb["Managed"].tables[anchors["table_name"]]
    broken_bytes = io.BytesIO()
    wb.save(broken_bytes)
    with pytest.raises(hashes.ObservedIdentityDriftError):
        hashes.compute_structure_hash_from_artifact(data=broken_bytes.getvalue(), contract=contract, anchors=anchors)
    for key in ("contract", "structure_anchors"):
        broken = copy.copy(plan)
        setattr(broken, key, None)
        with pytest.raises(cm.ContentMutationError):
            service._projection_structure_hash(plan=broken, output=output, materialized=legacy)
    plan.document_type = "docx"
    plan.structure_anchors = None
    assert service._projection_structure_hash(plan=plan, output=output, materialized=legacy) == legacy.structure_hash


@pytest.mark.asyncio
async def test_frozen_payload_reader_and_missing_anchors(monkeypatch):
    data, anchors = workbook_fixture()
    bundle = NS(bundle_id="frozen", authority_model=NS(value="projection_contract"))
    child = NS(kind="instrumentation", state="approved")
    read_child = AsyncMock(return_value=child)
    payload = {"managed_sheets": [{"sheet_key": anchors["sheet_key"],
        "region_boundary_locator": {"table_key": anchors["table_name"]},
        "tables": [{"row_uuid_column_letter": anchors["uuid_column_letter"]}]}],
        "hidden_metadata_sheet": {"sheet_name": anchors["metadata_sheet"]}}
    read_payload = AsyncMock(return_value=payload)
    monkeypatch.setattr(PublishedIdentityObserver, "_load_child_row", read_child)
    monkeypatch.setattr(PublishedIdentityObserver, "_read_definition_payload", read_payload)
    kwargs = dict(session=object(), resolution=object(), bundle=bundle, document_type="xlsx")
    assert await hashes.load_frozen_structure_anchors(**kwargs) == anchors
    assert read_child.call_args.kwargs["bundle"] is bundle
    read_payload.return_value = {}
    with pytest.raises(hashes.FrozenChildUnusableError):
        await hashes.load_frozen_structure_anchors(**kwargs)
    read_payload.return_value = payload
    child.state = "draft"
    with pytest.raises(hashes.FrozenChildUnusableError):
        await hashes.load_frozen_structure_anchors(**kwargs)
    read_child.reset_mock()
    assert await hashes.load_frozen_structure_anchors(**{**kwargs, "document_type": "docx"}) is None
    bundle.authority_model.value = "custom_authoritative_ooxml"
    assert await hashes.load_frozen_structure_anchors(**kwargs) is None
    read_child.assert_not_called()


@pytest.mark.parametrize("filename,expected", [
    ("materialize_coordinator.py", "pre.bundle"),
    ("oo_to_html.py", "state.bundle"),
    ("conflict_resolution.py", "resolved.bundle"),
])
@pytest.mark.asyncio
async def test_three_host_frozen_argument_execution_and_mutation(filename, expected):
    tree = ast.parse((ROOT / "backend/app/services/workpaper_sync" / filename).read_text(encoding="utf-8"))
    calls = [n for n in ast.walk(tree) if isinstance(n, ast.Call)
             and isinstance(n.func, ast.Name) and n.func.id == "ContentCommitPlan"]
    assert len(calls) == 1
    expressions = {k.arg: k.value for k in calls[0].keywords}
    expr = expressions["structure_anchors"]
    assert isinstance(expr, ast.Await)
    bundle = object()
    anchors = {"sentinel": "frozen-not-current"}
    seen = []
    async def load(**kwargs):
        seen.append(kwargs)
        assert kwargs["bundle"] is bundle
        return anchors
    namespace = dict(load_frozen_structure_anchors=load,
                     self=NS(_session=object(), _resolution=object()),
                     pre=NS(bundle=bundle, resolution=NS(document_type="xlsx")),
                     state=NS(bundle=bundle, substrate=NS(document_type="xlsx")),
                     resolved=NS(bundle=bundle, document_type="xlsx"))
    assert ast.unparse(next(k.value for k in expr.value.keywords if k.arg == "bundle")) == expected
    async def evaluate(expression):
        fn = ast.AsyncFunctionDef(name="probe", args=ast.arguments(posonlyargs=[], args=[], kwonlyargs=[], kw_defaults=[], defaults=[]),
                                 body=[ast.Return(expression)], decorator_list=[])
        module = ast.fix_missing_locations(ast.Module(body=[fn], type_ignores=[]))
        exec(compile(module, filename, "exec"), namespace)
        return await namespace["probe"]()
    assert await evaluate(expr) is anchors
    assert len(seen) == 1
    # Removing the actual host argument in memory must violate the same assertion.
    with pytest.raises(AssertionError):
        assert await evaluate(ast.Constant(None)) is anchors


def test_fail_open_mutation_is_killed(tmp_path):
    source = (ROOT / "backend/app/services/workpaper_sync/content_mutation.py").read_text(encoding="utf-8")
    tree = ast.parse(source)
    fn = next(n for n in ast.walk(tree) if isinstance(n, ast.FunctionDef) and n.name == "_projection_structure_hash")
    mutated = copy.deepcopy(fn)
    raises = [n for n in ast.walk(mutated) if isinstance(n, ast.If) and any(isinstance(b, ast.Raise) for b in n.body)]
    assert len(raises) == 1
    raises[0].body = [ast.Return(ast.Constant("legacy"))]
    module = ast.fix_missing_locations(ast.Module(body=[mutated], type_ignores=[]))
    ns = {"ContentCommitPlan": object, "Path": Path, "Any": object}
    exec(compile(module, "in-memory-fail-open", "exec"), ns)
    with pytest.raises(pytest.fail.Exception):
        with pytest.raises(cm.ContentMutationError):
            ns[fn.name](None, plan=NS(document_type="xlsx", structure_anchors=None, contract=None), output=tmp_path, materialized=NS())


@pytest.mark.asyncio
async def test_finalize_final_bytes_and_inline_formula_mutation(monkeypatch, tmp_path):
    from backend.tests.workpaper_sync import test_task36_excel_entry_gate as fixtures
    import app.services.workpaper_sync.excel_entry_gate as gate_module
    monkeypatch.setattr(gate_module, "load_contract", lambda _: contracts.parse_contract(fixtures.xlsx_payload()))
    from app.services.workpaper_sync.excel_entry_gate import ExcelEntryFinalizeGate
    from dataclasses import replace

    contract = contracts.parse_contract(fixtures.xlsx_payload())
    data, anchors = workbook_fixture()
    digest = hashlib.sha256(data).hexdigest()
    report = fixtures.evidence_bytes(instrumented_sha=digest)
    gate, coordinator, _ = fixtures.build_gate(contract.canonical_sha256, report=report)
    candidate = fixtures.staged_candidate(digest, tmp=tmp_path)
    candidate.path.write_bytes(data)
    frozen = AsyncMock(return_value=anchors)
    monkeypatch.setattr(hashes, "load_frozen_structure_anchors", frozen)
    expected = hashes.compute_structure_hash_from_artifact(data=data, contract=contract, anchors=anchors)
    kwargs = fixtures.gate_kwargs(report, sha=digest, staged_candidate=candidate)
    await gate.finalize_candidate(**kwargs)
    assert coordinator.calls[-1]["structure_hash"] == expected
    assert frozen.call_args.kwargs["bundle"].bundle_id == fixtures.BUNDLE_ID

    tree = ast.parse((ROOT / "backend/app/services/workpaper_sync/excel_entry_gate.py").read_text(encoding="utf-8"))
    fn = next(n for n in ast.walk(tree) if isinstance(n, ast.AsyncFunctionDef) and n.name == "finalize_candidate")
    mutated = copy.deepcopy(fn)
    assignments = [n for n in ast.walk(mutated) if isinstance(n, ast.Assign)
                   and any(isinstance(t, ast.Name) and t.id == "structure_hash" for t in n.targets)]
    assert len(assignments) == 1
    assignments[0].value = ast.parse("canonical_digest({'schema_version': 'excel-entry-structure:v1', 'contract_sha256': definitions.contract.canonical_sha256, 'structure': [['mutated', 'structure', 'coordinate', 'identity']]})", mode="eval").body
    import app.services.workpaper_sync.excel_entry_gate as module
    namespace = dict(vars(module))
    exec(compile(ast.fix_missing_locations(ast.Module(body=[mutated], type_ignores=[])), "in-memory-finalize", "exec"), namespace)
    await namespace[fn.name](gate, **kwargs)
    with pytest.raises(AssertionError):
        assert coordinator.calls[-1]["structure_hash"] == expected
    # Changed bytes must stop before the coordinator even when metadata still agrees.
    before = len(coordinator.calls)
    candidate.path.write_bytes(b"tampered")
    with pytest.raises(module.CandidateEvidenceError):
        await gate.finalize_candidate(**kwargs)
    assert len(coordinator.calls) == before


@pytest.mark.asyncio
async def test_frozen_child_and_payload_digest_fail_closed(monkeypatch, tmp_path):
    import json
    import uuid
    from app.services.workpaper_sync.definitions import BundleSlot, canonical_digest
    from app.services.workpaper_sync.published_identity_observer import FrozenChildUnusableError

    payload = {"managed_sheets": []}
    child = NS(id=uuid.uuid4(), sha256=canonical_digest(payload))
    slot = NS(is_definition=True, slot_ref=f"definition:{child.id}", slot_digest=child.sha256)
    bundle = NS(slots={BundleSlot.instrumentation: slot})
    session = NS(execute=AsyncMock(return_value=NS(scalar_one_or_none=lambda: child)))
    observer = PublishedIdentityObserver(session=session, resolution=object())
    assert await observer._load_child_row(slot=BundleSlot.instrumentation, bundle=bundle, context={}) is child
    slot.slot_digest = "wrong"
    with pytest.raises(FrozenChildUnusableError):
        await observer._load_child_row(slot=BundleSlot.instrumentation, bundle=bundle, context={})
    artifact = tmp_path / "definition.json"
    artifact.write_text(json.dumps(payload), encoding="utf-8")
    monkeypatch.setattr(observer, "_resolve_definition_blob", AsyncMock(return_value=artifact))
    assert await observer._read_definition_payload(child=child, context={}) == payload
    artifact.write_text('{"tampered": true}', encoding="utf-8")
    with pytest.raises(FrozenChildUnusableError):
        await observer._read_definition_payload(child=child, context={})

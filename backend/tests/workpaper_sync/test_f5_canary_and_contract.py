# -*- coding: utf-8 -*-
"""F5 canary 判据：Property 1 / 2 / 3 / 4 / 6 / 17 / 19 / 24 + 锚行 footer + 契约可解析。

spec: f5-sync-coverage-and-first-canary
      Task 3（P2/P4/P17/P24）· Task 5（P22 基线）· Task 7（provider）· Task 8（canary 薄声明）
      Requirements 1.1 / 1.2 / 1.3 / 1.4 / 1.6 / 2.1 / 2.2 / 2.4 / 2.5 / 6.1 / 7.3 / 7.5

🔴 Property 编号 spec-scoped：本文件的 `Property N` 读作 `F5-P{N}`。

F5 的四处特殊性都在本文件有判据：
  ① 真库**完全无载荷** ⇒ 裁决条目 `max_payload_bytes == 0`，验收必须先 seed（P2④ / P8）
  ② **无 footer 合计** ⇒ 锚行 + `carries_total_formula=False`（P3）
  ③ 行身份 **`id`** 不是 `rowId`（P4）
  ④ UUID 列 **超出 `max_column`** ⇒ instrumentation 需扩列（P6）
"""
from __future__ import annotations

import json
import sys
import unittest.mock as mock
from pathlib import Path

import pytest

_BACKEND = Path(__file__).resolve().parents[2]
if str(_BACKEND) not in sys.path:
    sys.path.insert(0, str(_BACKEND))

from app.services.workpaper_sync import phase5_f5_08_major_adjustment as F508  # noqa: E402
from app.services.workpaper_sync import phase5_f5_cost_of_sales as F5  # noqa: E402
from app.services.workpaper_sync.contracts import parse_contract  # noqa: E402
from app.services.workpaper_sync.entry_profile import (  # noqa: E402
    load_entry_manifest,
    manifest_entries_by_id,
)
from app.services.workpaper_sync.excel_extract import BindingKind  # noqa: E402
from app.services.workpaper_sync.phase5_row_table_sheet import StoreKind  # noqa: E402

F5_ENTRY_ID = "xlsx/gt-f5-cost-of-sales"


@pytest.fixture(scope="module")
def manifest_entry() -> dict:
    return manifest_entries_by_id(load_entry_manifest())[F5_ENTRY_ID]


@pytest.fixture(scope="module")
def contract_payload() -> dict:
    return F5.build_contract_payload()


@pytest.fixture(scope="module")
def slice_entry() -> dict:
    data = json.loads(
        (_BACKEND / "data" / "workpaper_sync_f_cycle_manifest_slice.json").read_text(
            encoding="utf-8"
        )
    )
    return next(e for e in data["independent_entries"] if e["entry_id"] == F5_ENTRY_ID)


def _resolution(codes) -> F5.TemplateResolutionFacts:
    return F5.TemplateResolutionFacts(
        by_wp_code={c: [None] for c in codes},
        parent_code="F5",
        parent_resolved_path=F5.authoritative_template_path(),
    )


# ═══════════════════════════════════════════════════════════════════════════
# F5-P1：migration_state（🔴 现状必红）+ slice 逐元素（含 BP-7 不含 BP-5）
# ═══════════════════════════════════════════════════════════════════════════


class TestProperty1AndSlice:
    def test_current_state_is_legacy(self, manifest_entry) -> None:
        assert manifest_entry.get("migration_state") == "legacy_fake_bidirectional"

    @pytest.mark.xfail(
        reason="🔴 现状必红：published representation 缺供给（BP-61-1），Task 9 就绪后转绿",
        strict=True,
    )
    def test_becomes_adapter_registered(self, manifest_entry) -> None:
        assert manifest_entry.get("migration_state") == "adapter_registered"

    def test_slice_blocked_by_includes_bp7_excludes_bp5(self, slice_entry) -> None:
        """🔴 Task 1 要求的逐元素断言：F5 含 BP-7，**不含 BP-5**（BP-5 是 F2 专属）。"""
        blocked = slice_entry["capability_target_blocked_by"]
        assert blocked == ["BP-1", "BP-2", "BP-3", "BP-4", "BP-7"]
        assert "BP-5" not in blocked
        assert "BP-9" not in blocked

    def test_slice_five_supply_slots_are_null(self, slice_entry) -> None:
        """BP-61-1 属实的直接证据。"""
        for key in (
            "adapter_id",
            "authority_model",
            "definition_bundle",
            "instrumentation_candidate",
            "published_representation",
        ):
            assert slice_entry[key] is None, f"{key} 应为 null（供给缺口）"

    def test_slice_capability_diverges_from_manifest(self, slice_entry, manifest_entry) -> None:
        """FC-12：slice 的 `capability=null` 与 manifest 的 overlay 组件级默认值不一致且已登记。"""
        assert slice_entry["capability"] is None
        assert slice_entry["capability_target"] == "bidirectional"
        assert slice_entry["manifest_mirror"]["capability"] == "single_onlyoffice"
        assert slice_entry["manifest_mirror"]["divergence_from_slice"]


# ═══════════════════════════════════════════════════════════════════════════
# F5-P2：幻影码隔离 + 🔴 零载荷证据如实记 0
# Validates: 1.2 / 1.3
# ═══════════════════════════════════════════════════════════════════════════


class TestProperty2PhantomAndZeroPayload:
    def test_wp_codes_is_phantom(self) -> None:
        assert F5.WP_CODES == frozenset({"F5C"})

    def test_index_json_has_f5_but_not_f5c(self) -> None:
        index = json.loads((_BACKEND / "wp_templates" / "_index.json").read_text(encoding="utf-8"))
        codes = {f.get("wp_code") for f in index["files"]}
        assert "F5" in codes and "F5C" not in codes

    def test_adjudication_records_zero_payload_honestly(self) -> None:
        """🔴 需求 1.3：F5 全部键真库 0 行 ⇒ `max_payload_bytes` 必须如实记 0，不得伪造。"""
        adj = json.loads(
            (_BACKEND / "data" / "workpaper_sync_entry_wp_code_adjudication.json").read_text(
                encoding="utf-8"
            )
        )
        entry = next(e for e in adj["adjudications"] if e["entry_id"] == F5_ENTRY_ID)
        assert entry["wp_codes"] == ["F5"]
        evidence = entry["store_payload_evidence"]
        assert evidence["max_payload_bytes"] == 0
        assert evidence["wp_count_with_payload"] == 0
        assert evidence["wp_code_with_payload"] is None
        assert "0 行" in evidence["note"], "note 须点明零载荷事实"
        assert "seed" in evidence["note"], "note 须点明验收前必须 seed（裁决 F5-H7）"

    def test_fallback_guard_passes(self) -> None:
        F5.assert_no_implicit_template_fallback(
            _resolution(F5.WP_CODES), wp_codes=frozenset({"F5C"})
        )

    def test_mutation_true_code_raises(self) -> None:
        with mock.patch.object(F5, "WP_CODES", frozenset({"F5"})):
            with pytest.raises(F5.EntrySelectionError, match="wp_code_patterns"):
                F5.assert_entry_selectable(resolution=_resolution({"F5"}))


# ═══════════════════════════════════════════════════════════════════════════
# F5-P3：无 footer 合计 ⇒ 锚行 + carries_total_formula=False
# Validates: 1.4 / 2.4
# ═══════════════════════════════════════════════════════════════════════════


class TestProperty3AnchorFooter:
    def test_footer_carries_no_total_formula(self) -> None:
        assert F508.SPEC_F508.footer_carries_total_formula is False

    def test_footer_marker_matches_template_verbatim(self) -> None:
        """🔴 逐字实测 A30 带**全角冒号** —— spec Task 8 写的「三、审计说明」漏了冒号，
        `assert_footer_anchor_stable` 是逐字匹配，用 spec 的值会定位失败。
        """
        from openpyxl import load_workbook

        wb = load_workbook(F5.authoritative_template_path(), data_only=False)
        try:
            ws = wb[F508.MANAGED_SHEET_F508]
            actual = ws.cell(row=F508.FOOTER_ROW_F508, column=1).value
        finally:
            wb.close()
        assert actual == "三、审计说明："
        assert F508.FOOTER_MARKER_F508 == actual
        assert F508.FOOTER_MARKER_F508.endswith("："), "全角冒号必须保留"

    def test_anchor_row_has_no_formula(self) -> None:
        """锚行确实没有合计公式（否则不该声明 False）。"""
        from openpyxl import load_workbook

        wb = load_workbook(F5.authoritative_template_path(), data_only=False)
        try:
            ws = wb[F508.MANAGED_SHEET_F508]
            row_values = [
                ws.cell(row=F508.FOOTER_ROW_F508, column=ci).value
                for ci in range(1, ws.max_column + 1)
            ]
        finally:
            wb.close()
        formulas = [v for v in row_values if isinstance(v, str) and v.startswith("=")]
        assert formulas == [], f"锚行 R30 不该有公式，实得 {formulas}"

    def test_contract_carries_flag_propagates(self, contract_payload) -> None:
        footer = contract_payload["sheets"][0]["tables"][0]["footer_anchor"]
        assert footer["carries_total_formula"] is False
        assert footer["marker"] == "三、审计说明："

    def test_html_only_anchor_rows_registered(self) -> None:
        """锚行之后的审计说明/结论/提示区登记为 HTML-only（不是漏声明）。"""
        rows = dict(F508.HTML_ONLY_ANCHOR_ROWS_F508)
        assert rows[30] == "三、审计说明："
        assert rows[33] == "四、审计结论："
        assert 37 in rows and 38 in rows and 39 in rows


# ═══════════════════════════════════════════════════════════════════════════
# F5-P4：行身份是 `id` 不是 `rowId` + 主键不取 legacy
# Validates: 2.1 / 2.2
# ═══════════════════════════════════════════════════════════════════════════


class TestProperty4RowIdentityIsId:
    def test_row_identity_key_is_id(self) -> None:
        """🔴 F5-2/3/5/8 四张皆用 `id`，与 D 类惯例相反（同 E1 教训）。"""
        assert F508.SPEC_F508.row_identity_key == "id"
        assert F5.ROW_IDENTITY_STORE_KEY == "id"

    def test_contract_row_identity_pointer_uses_id(self, contract_payload) -> None:
        table = contract_payload["sheets"][0]["tables"][0]
        assert table["row_identity"] == {"kind": "field", "json_pointer": "/rows/*/id"}
        assert "/rows/*/rowId" != table["row_identity"]["json_pointer"]

    def test_triad_is_excel_table(self) -> None:
        assert F508.SPEC_F508.binding_kind is BindingKind.excel_table
        assert F508.SPEC_F508.store_kind is StoreKind.rows

    @pytest.mark.parametrize(
        "mutated",
        ["F5-8-conclusion", "F5-2-rows", "F5-7-rows", "F5-2-monthly-rows"],
    )
    def test_non_managed_keys_raise(self, mutated) -> None:
        """变异：legacy 读回退键 / 不存在的键 / 未启用区的键 ⇒ 必抛。"""
        with pytest.raises(F5.StorePayloadError):
            F5._spec_of_store_item(mutated)


# ═══════════════════════════════════════════════════════════════════════════
# F5-P6：UUID 列超出 max_column ⇒ instrumentation 需扩列
# Validates: 2.5
# ═══════════════════════════════════════════════════════════════════════════


class TestProperty6UuidColumnBeyondMaxCol:
    def test_uuid_column_exceeds_template_max_column(self) -> None:
        """🔴 F5-8 的 UUID 列 I 超出模板 max_column(H) ⇒ 扩列场景（需求 2.5）。"""
        from openpyxl import load_workbook
        from openpyxl.utils import column_index_from_string

        wb = load_workbook(F5.authoritative_template_path(), data_only=False)
        try:
            ws = wb[F508.MANAGED_SHEET_F508]
            max_col = ws.max_column
        finally:
            wb.close()
        assert max_col == 8, f"F5-8 模板 max_column 实测应为 8(H)，实得 {max_col}"
        assert column_index_from_string(F508.UUID_COL_F508) == 9, "UUID 列 I = 第 9 列"
        assert column_index_from_string(F508.UUID_COL_F508) > max_col

    def test_uuid_column_is_empty_in_template(self) -> None:
        from openpyxl import load_workbook

        wb = load_workbook(F5.authoritative_template_path(), data_only=False)
        try:
            ws = wb[F508.MANAGED_SHEET_F508]
            values = [ws.cell(row=r, column=9).value for r in range(1, ws.max_row + 1)]
        finally:
            wb.close()
        assert all(v is None for v in values), "UUID 列必须在模板中全空"


# ═══════════════════════════════════════════════════════════════════════════
# F5-P17 / P24：HTML-only 与零写入读键登记而不误接
# Validates: 6.1 / 7.5
# ═══════════════════════════════════════════════════════════════════════════


class TestProperty17And24Registrations:
    def test_f51_main_zone_is_html_only(self) -> None:
        """🔴 裁决 F5-H3：F5-1 主营区零 editable 字段 ⇒ 不进受管清单（否则与 F5-2 双源）。"""
        keys = dict(F5.HTML_ONLY_STORE_KEYS)
        assert "F5-1-adj-main-rows" in keys
        assert "F5-2" in keys["F5-1-adj-main-rows"], "登记须说明与 F5-2 的双源风险"
        assert "F5-1-adj-main-rows" not in F5.all_store_item_ids()

    def test_f54_hub_and_f56_are_html_only(self) -> None:
        keys = dict(F5.HTML_ONLY_STORE_KEYS)
        assert "F5-4-rows" in keys and "FC-6" in keys["F5-4-rows"]
        assert "F5-6-quantity-recon-rows" in keys
        for key in keys:
            assert key not in F5.all_store_item_ids()

    def test_zero_writer_read_keys_registered(self) -> None:
        """P24：两个零写入读键登记为兼容回退，不进 field_specs。"""
        assert F5.ZERO_WRITER_READ_KEYS == ("D4-1-adj-main-rows", "F5-2-detail-rows")
        for key in F5.ZERO_WRITER_READ_KEYS:
            assert key not in F5.all_store_item_ids()

    def test_cross_sheet_injected_key_registered(self) -> None:
        """P16：`F5-7-adjudicated-cogs` 由宿主注入（非本表数据）⇒ 不进 field_specs。"""
        injected = dict(F5.CROSS_SHEET_INJECTED_KEYS)
        assert "F5-7-adjudicated-cogs" in injected
        assert "substantive:adjudicated" in injected["F5-7-adjudicated-cogs"]
        assert "F5-7-adjudicated-cogs" not in F5.all_store_item_ids()


# ═══════════════════════════════════════════════════════════════════════════
# F5-P19（部分）：FC-10 在 F5 canary 不命中（逐列取证）
# Validates: 7.3
# ═══════════════════════════════════════════════════════════════════════════


class TestProperty19Fc10NotApplicableOnCanary:
    def test_canary_sheet_zero_percent_format(self) -> None:
        from openpyxl import load_workbook

        wb = load_workbook(F5.authoritative_template_path(), data_only=False)
        try:
            ws = wb[F508.MANAGED_SHEET_F508]
            pct = [
                f"{c.column_letter}{c.row}"
                for row in ws.iter_rows()
                for c in row
                if "%" in (c.number_format or "")
            ]
        finally:
            wb.close()
        assert pct == [], f"F5-8 实测应零百分比格式格，实得 {pct}"


# ═══════════════════════════════════════════════════════════════════════════
# 契约装配 + 选型真调
# ═══════════════════════════════════════════════════════════════════════════


class TestContractAndSelection:
    def test_payload_passes_parse_contract(self, contract_payload) -> None:
        contract = parse_contract(contract_payload, adapter_id=F5.ADAPTER_ID)
        assert contract.contract_id == "f5.cost_of_sales_detail"

    def test_table_structural_keys(self, contract_payload) -> None:
        table = contract_payload["sheets"][0]["tables"][0]
        assert table["anchor"] == "A12", "两级表头 ⇒ anchor 取 header_group_row"
        assert table["header_rows"] == 2
        assert table["uuid_col"] == "I"
        assert table["formula_mask"] == [], "数据区零公式"
        assert len(table["fields"]) == 7, "8 列受管区 - G 列（被 F 合并）= 7"

    def test_g_column_not_declared(self, contract_payload) -> None:
        """G 列被 F 的 F{r}:G{r} 合并吞掉 ⇒ 不单独声明字段。"""
        table = contract_payload["sheets"][0]["tables"][0]
        columns = [f["cell"]["column"] for f in table["fields"]]
        assert columns == ["A", "B", "C", "D", "E", "F", "H"]
        assert "G" not in columns

    def test_merged_cells_justify_skipping_g(self) -> None:
        """逐格证明 G 被合并（不是漏声明）。"""
        from openpyxl import load_workbook

        wb = load_workbook(F5.authoritative_template_path(), data_only=False)
        try:
            ws = wb[F508.MANAGED_SHEET_F508]
            merged = {str(m) for m in ws.merged_cells.ranges}
        finally:
            wb.close()
        assert "F12:G13" in merged, "表头区 F/G 跨两行两列合并"
        for r in range(F508.FIRST_DATA_ROW_F508, F508.LAST_DATA_ROW_F508 + 1):
            assert f"F{r}:G{r}" in merged, f"数据行 R{r} 的 F:G 应合并"

    def test_disk_contract_matches_source(self) -> None:
        assert F5.assert_contract_file_matches_source().contract_id == "f5.cost_of_sales_detail"

    def test_real_call_on_real_manifest(self) -> None:
        entry = F5.assert_entry_selectable(resolution=_resolution(F5.WP_CODES))
        assert entry["entry_id"] == F5_ENTRY_ID

    def test_build_matcher_constructible(self) -> None:
        matcher = F5.build_matcher()
        assert matcher.document_type == "xlsx"
        assert set(matcher.wp_codes) == {"F5C"}

    def test_projection_uses_id_as_row_identity(self) -> None:
        """零载荷 entry 的投影仍可用合成载荷验证键形态（行身份走 `id`）。"""
        contract = F5.load_contract_from_disk()
        payload = json.dumps(
            [
                {
                    "id": "f5maj-1",
                    "date": "2025-06-30",
                    "voucherNo": "JZ-001",
                    "itemContent": "返利冲减成本",
                    "debitAmount": 0,
                    "creditAmount": 12000,
                    "netAmount": -12000,
                    "adjustmentReason": "供应商年度返利",
                    "reasonAdequate": "是",
                }
            ],
            ensure_ascii=False,
        )
        projection = F5.build_store_projection(payload, contract=contract)
        assert len(projection.values) == 7
        assert "major_adjustment_rows/f5maj-1/credit_amount" in projection.values
        assert projection.values["major_adjustment_rows/f5maj-1/credit_amount"].value == 12000

    def test_projection_rejects_row_missing_id(self) -> None:
        """fail-closed：缺 `id`（而非缺 `rowId`）时必抛 —— 行身份键是 `id`。"""
        from app.services.workpaper_sync.phase5_row_table_sheet import (
            RowTableStorePayloadError,
        )

        contract = F5.load_contract_from_disk()
        bad = json.dumps([{"rowId": "wrong-key", "voucherNo": "X"}], ensure_ascii=False)
        with pytest.raises(RowTableStorePayloadError, match="稳定行身份"):
            F5.build_store_projection(bad, contract=contract)


# ══════════════════════════════════════════════════════════════════════════════
# Task 21：F5-1 审定表「其他业务成本」区
#
# 证据: .kiro/specs/f5-sync-coverage-and-first-canary/evidence/task21-adjudication-f5-1-geometry.md
# 四个与 F5-8 形成**对照**的判定（同册两张表在这四点上取值相反，照抄任一张都会错）：
#   ① 行身份 `rowKey` ↔ F5-8 的 `id`
#   ② uuid_col K **在** max_column 内 ↔ F5-8 的 I 超出
#   ③ footer 携带合计公式 True ↔ F5-8 的锚行 False
#   ④ 有公式列 (E,I) ↔ F5-8 的零公式 ()
# ══════════════════════════════════════════════════════════════════════════════
from app.services.workpaper_sync import phase5_f5_01_adjudication as F501  # noqa: E402


def _strip_ts_comments(text: str) -> str:
    """剥掉 TS 的块注释与行注释，只留活代码。

    行注释判定排除 `https://` 这类协议分隔符（`//` 前一字符是 `:` 时不算注释）。
    """
    import re

    without_block = re.sub(r"/\*.*?\*/", "", text, flags=re.S)
    return re.sub(r"(?<!:)//[^\n]*", "", without_block)


def _f501_table(payload: dict) -> dict:
    """从契约 payload 取 F5-1 其他业务区的 table（按 sheet_key/table_key 找，不用下标）。"""
    for sheet in payload["sheets"]:
        if sheet["sheet_key"] != F501.SHEET_KEY_F501:
            continue
        for table in sheet["tables"]:
            if table["table_key"] == F501.ROWS_TABLE_KEY_F501_OTHER:
                return table
    raise AssertionError(
        f"契约缺 {F501.SHEET_KEY_F501}/{F501.ROWS_TABLE_KEY_F501_OTHER}；"
        f"实得 sheets={[s['sheet_key'] for s in payload['sheets']]}"
    )


class TestF501MainZoneIsUnmanageable:
    """🔴 主营区不是"先不接"，是**不该接**：九列全公式、零可编辑格。"""

    def test_main_zone_registered_as_unmanageable(self) -> None:
        rows, reason = F501.UNMANAGED_MAIN_ZONE_F501
        assert rows == "R8:R17"
        assert "零用户输入点" in reason
        assert "F5-1-adj-main-rows" in reason, "须点明对应 store key，便于反查"

    def test_main_zone_store_key_stays_out_of_managed_ids(self) -> None:
        """加了 other 区之后，主营区 key 仍不得进受管集合（双源防线）。"""
        assert "F5-1-adj-main-rows" not in F5.all_store_item_ids()
        assert F501.STORE_ITEM_ID_F501_OTHER in F5.all_store_item_ids()

    def test_main_zone_nine_columns_are_all_formulas(self) -> None:
        """逐格证明"不该接"：R8~R17 的 A~I 九列每格都是公式 ⇒ 无处可写。"""
        from openpyxl import load_workbook

        wb = load_workbook(
            _BACKEND / "wp_templates" / "F" / "F5 营业成本.xlsx", data_only=False
        )
        ws = wb[F501.MANAGED_SHEET_F501]
        try:
            for row in range(8, 18):
                for col in "ABCDEFGHI":
                    value = ws[f"{col}{row}"].value
                    assert isinstance(value, str) and value.startswith("="), (
                        f"{col}{row} 应是公式（主营区九列全公式），实得 {value!r}"
                    )
        finally:
            wb.close()

    def test_other_zone_has_editable_cells(self) -> None:
        """对照：其他业务区 R20~R25 只有 E/I 是公式，其余列可编辑。"""
        from openpyxl import load_workbook

        wb = load_workbook(
            _BACKEND / "wp_templates" / "F" / "F5 营业成本.xlsx", data_only=False
        )
        ws = wb[F501.MANAGED_SHEET_F501]
        try:
            for row in range(20, 26):
                for col in ("E", "I"):
                    value = ws[f"{col}{row}"].value
                    assert isinstance(value, str) and value.startswith("=")
                for col in ("A", "B", "C", "D", "F", "G", "H", "J"):
                    value = ws[f"{col}{row}"].value
                    assert value is None or not (
                        isinstance(value, str) and value.startswith("=")
                    ), f"{col}{row} 应可编辑（无公式），实得 {value!r}"
        finally:
            wb.close()


class TestF501RowIdentityIsRowKey:
    """① 行身份 `rowKey` —— 与同册 F5-8 的 `id` 相反（FC-4：按值 grep，不推演）。"""

    def test_row_identity_key_is_row_key_not_id(self) -> None:
        assert F501.SPEC_F501_OTHER.row_identity_key == "rowKey"
        assert F501.ROW_IDENTITY_STORE_KEY_F501 == "rowKey"

    def test_diverges_from_f508_in_same_workbook(self) -> None:
        """🔴 同一册内两种行身份键 —— 照抄 F5-8 会让投影全行 fail-closed。"""
        assert F508.SPEC_F508.row_identity_key == "id"
        assert F501.SPEC_F501_OTHER.row_identity_key != F508.SPEC_F508.row_identity_key

    def test_contract_pointer_uses_row_key(self, contract_payload) -> None:
        table = _f501_table(contract_payload)
        assert table["row_identity"] == {
            "kind": "field",
            "json_pointer": "/rows/*/rowKey",
        }

    def test_store_item_id_matches_frontend_constant(self) -> None:
        """按值实测：`useF5Adjudication.ts` 的 `OTHER_STORAGE_KEY`。"""
        assert F501.STORE_ITEM_ID_F501_OTHER == "F5-1-adj-other-rows"

    def test_projection_rejects_row_missing_row_key(self) -> None:
        """fail-closed：缺 `rowKey` 必抛，绝不退回数组下标。"""
        from app.services.workpaper_sync.phase5_row_table_sheet import (
            RowTableStorePayloadError,
            build_store_projection,
        )

        contract = F5.load_contract_from_disk()
        with pytest.raises(RowTableStorePayloadError):
            build_store_projection(
                spec=F501.SPEC_F501_OTHER,
                contract=contract,
                payload=json.dumps([{"label": "运输成本", "currentUnadjusted": 100}]),
            )

    def test_projection_roundtrip_with_row_key(self) -> None:
        """带 rowKey 的合成载荷可投影（真库 0 行，故用合成载荷验键形态）。"""
        from app.services.workpaper_sync.phase5_row_table_sheet import (
            build_store_projection,
        )

        contract = F5.load_contract_from_disk()
        projection = build_store_projection(
            spec=F501.SPEC_F501_OTHER,
            contract=contract,
            payload=json.dumps(
                [
                    {
                        "rowKey": "other-1",
                        "label": "运输成本",
                        "isFixed": False,
                        "currentUnadjusted": 1000,
                        "currentAje": 0,
                        "currentRje": 0,
                        "priorUnadjusted": 900,
                        "priorAje": 0,
                        "priorRje": 0,
                        "indexRef": "F5-3",
                    }
                ]
            ),
        )
        key = f"{F501.ROWS_TABLE_KEY_F501_OTHER}/other-1/current_unadjusted"
        assert projection.values[key].value == 1000
        prior = f"{F501.ROWS_TABLE_KEY_F501_OTHER}/other-1/prior_unadjusted"
        assert projection.values[prior].value == 900

    def test_store_only_key_is_registered(self) -> None:
        """`isFixed` 模板无列 ⇒ 登记为 store-only，不进 field_specs。"""
        store_only = dict(F501.STORE_ONLY_KEYS_F501_OTHER)
        assert "isFixed" in store_only
        json_keys = {spec[4] for spec in F501.FIELD_SPECS_F501_OTHER}
        assert "isFixed" not in json_keys


class TestF501UuidColumnWithinMaxCol:
    """② uuid_col K **在** max_column 内 —— 与 F5-8 的 I 超出形成对照（扩列与否的分野）。"""

    def test_uuid_col_is_k_not_j(self) -> None:
        """🔴 修正 spec：spec design 记「J~N 全空」有误，J 有值（J5:J6 索引表头 + J3 公式）。"""
        assert F501.UUID_COL_F501_OTHER == "K"

    def test_column_j_is_not_empty_in_template(self) -> None:
        """逐格证明 J 不可用作 uuid_col（否则会覆盖索引表头）。"""
        from openpyxl import load_workbook

        wb = load_workbook(
            _BACKEND / "wp_templates" / "F" / "F5 营业成本.xlsx", data_only=False
        )
        ws = wb[F501.MANAGED_SHEET_F501]
        try:
            occupied = [
                f"J{r}" for r in range(1, ws.max_row + 1) if ws[f"J{r}"].value is not None
            ]
            assert occupied, "J 列应有非空格（证明它不是全空列）"
        finally:
            wb.close()

    def test_uuid_col_k_is_empty_in_template(self) -> None:
        from openpyxl import load_workbook

        wb = load_workbook(
            _BACKEND / "wp_templates" / "F" / "F5 营业成本.xlsx", data_only=False
        )
        ws = wb[F501.MANAGED_SHEET_F501]
        try:
            for r in range(1, ws.max_row + 1):
                cell = ws[f"{F501.UUID_COL_F501_OTHER}{r}"]
                assert cell.value is None, f"uuid 列 K{r} 应为空，实得 {cell.value!r}"
        finally:
            wb.close()

    def test_uuid_col_within_max_column_unlike_f508(self) -> None:
        """K 在 max_column(N) 内 ⇒ **无需扩列**；F5-8 的 I 超出 max_column(H) ⇒ 需扩列。"""
        from openpyxl import load_workbook
        from openpyxl.utils import column_index_from_string

        wb = load_workbook(
            _BACKEND / "wp_templates" / "F" / "F5 营业成本.xlsx", data_only=False
        )
        try:
            ws = wb[F501.MANAGED_SHEET_F501]
            assert column_index_from_string(F501.UUID_COL_F501_OTHER) <= ws.max_column

            ws508 = wb[F508.MANAGED_SHEET_F508]
            assert column_index_from_string(F508.UUID_COL_F508) > ws508.max_column, (
                "对照项：F5-8 的 uuid 列必须仍超出 max_column（否则扩列判据失去对照）"
            )
        finally:
            wb.close()

    def test_sibling_uuid_cols_are_distinct(self) -> None:
        """框架层用 uuid_col 配对 spec↔table ⇒ 本 entry 内所有受管区的 uuid_col 必须互不相同。"""
        specs = F5.managed_row_table_specs()
        cols = [s.uuid_col for s in specs]
        assert len(cols) == len(set(cols)), f"uuid_col 重复：{cols}"


class TestF501FooterCarriesTotalFormula:
    """③ footer R26 携带合计公式 True —— 与 F5-8 的锚行 False 形成对照。"""

    def test_flag_is_true_unlike_f508(self) -> None:
        assert F501.SPEC_F501_OTHER.footer_carries_total_formula is True
        assert F508.SPEC_F508.footer_carries_total_formula is False

    def test_footer_marker_is_subtotal(self) -> None:
        assert F501.FOOTER_MARKER_F501_OTHER == "小计"
        assert F501.SPEC_F501_OTHER.footer_row == 26

    def test_footer_row_really_has_sum_formulas(self) -> None:
        """逐格证明 R26 确有 SUM 公式（否则不该声明 True）。"""
        from openpyxl import load_workbook

        wb = load_workbook(
            _BACKEND / "wp_templates" / "F" / "F5 营业成本.xlsx", data_only=False
        )
        ws = wb[F501.MANAGED_SHEET_F501]
        try:
            assert ws["A26"].value == F501.FOOTER_MARKER_F501_OTHER
            for col in "BCDEFGHI":
                value = ws[f"{col}26"].value
                assert isinstance(value, str) and value.startswith("=SUM("), (
                    f"{col}26 应是 SUM 公式，实得 {value!r}"
                )
                assert "20:" in value and "25" in value, (
                    f"{col}26 的 SUM 区间应覆盖数据区 R20:R25，实得 {value}"
                )
        finally:
            wb.close()

    def test_contract_footer_anchor(self, contract_payload) -> None:
        footer = _f501_table(contract_payload)["footer_anchor"]
        assert footer["carries_total_formula"] is True
        assert footer["marker"] == "小计"


class TestF501FormulaColumnsAndFields:
    """④ 有公式列 (E,I) —— 与 F5-8 的零公式 () 形成对照；字段 8 个（E/I 不进 fields）。"""

    def test_formula_columns_are_e_and_i(self) -> None:
        assert F501.SPEC_F501_OTHER.formula_columns == ("E", "I")
        assert F508.SPEC_F508.formula_columns == (), "对照项：F5-8 数据区零公式"

    def test_formula_templates_match_template_verbatim(self) -> None:
        """模板逐字：E=B+C+D（本期审定）/ I=F+G+H（上期审定）。"""
        templates = F501.SPEC_F501_OTHER.formula_templates
        assert templates["E"] == "=B{r}+C{r}+D{r}"
        assert templates["I"] == "=F{r}+G{r}+H{r}"

    def test_templates_reproduce_every_data_row(self) -> None:
        """把声明的模板按行渲染，逐行与模板真实公式比对（六行全覆盖）。"""
        from openpyxl import load_workbook

        wb = load_workbook(
            _BACKEND / "wp_templates" / "F" / "F5 营业成本.xlsx", data_only=False
        )
        ws = wb[F501.MANAGED_SHEET_F501]
        try:
            for row in range(
                F501.FIRST_DATA_ROW_F501_OTHER, F501.LAST_DATA_ROW_F501_OTHER + 1
            ):
                for col, tpl in F501.SPEC_F501_OTHER.formula_templates.items():
                    assert ws[f"{col}{row}"].value == tpl.format(r=row)
        finally:
            wb.close()

    def test_field_count_and_columns(self, contract_payload) -> None:
        table = _f501_table(contract_payload)
        assert len(table["fields"]) == 8, "A,B,C,D,F,G,H,J 八列可编辑"
        columns = {f["cell"]["column"] for f in table["fields"]}
        assert columns == {"A", "B", "C", "D", "F", "G", "H", "J"}
        assert "E" not in columns and "I" not in columns, "公式列不进 fields"

    def test_anchor_uses_shared_two_level_header(self, contract_payload) -> None:
        """区② 无独立列表头 ⇒ 复用表级 R5/R6，anchor=A5（D3-4 双区先例允许跨度）。"""
        table = _f501_table(contract_payload)
        assert table["anchor"] == "A5"
        assert table["header_rows"] == 2
        assert F501.HEADER_GROUP_ROW_F501 == 5
        assert F501.HEADER_LEAF_ROW_F501 == 6

    def test_data_range_and_store_kind(self) -> None:
        spec = F501.SPEC_F501_OTHER
        assert (spec.first_data_row, spec.last_data_row) == (20, 25)
        assert spec.binding_kind is BindingKind.excel_table
        assert spec.store_kind is StoreKind.rows

    def test_sheet_key_is_single_shared_value(self) -> None:
        """🔴 D3-4 先例：同 managed_sheet 的多区必须共享单一 sheet_key（不是逐区起名）。"""
        assert F501.SHEET_KEY_F501 == "f51-managed"
        assert F501.SPEC_F501_OTHER.sheet_key == F501.SHEET_KEY_F501
        assert F501.SHEET_KEY_F501 != F508.SHEET_KEY_F508, "不同 sheet 必须不同 sheet_key"

    def test_table_name_and_template_id_are_zone_scoped(self) -> None:
        """区级唯一性靠 table_name/template_id（sheet_key 共享时的区分维度）。"""
        spec = F501.SPEC_F501_OTHER
        assert spec.table_name == "GT_F51_OTHER_ROWS"
        assert spec.template_id == "F51OTHER"
        names = {s.table_name for s in F5.managed_row_table_specs()}
        assert len(names) == len(F5.managed_row_table_specs()), f"table_name 重复：{names}"

    def test_html_only_rows_registered(self) -> None:
        """区①标题/两处小计/合计/试算/差异登记为 HTML-only（证明不是漏声明）。"""
        rows = dict(F501.HTML_ONLY_ANCHOR_ROWS_F501)
        assert rows[7] == "主营业业成本："
        assert rows[18] == "小计"
        assert rows[19] == "其他业务成本"
        assert rows[27] == "合计"
        assert rows[28] == "试算平衡表数"
        assert rows[29] == "差异数"

    def test_contract_still_parses_with_two_sheets(self, contract_payload) -> None:
        """加了 F5-1 之后整份契约仍过 parse_contract（两个 sheet 共存）。"""
        contract = parse_contract(contract_payload, adapter_id=F5.ADAPTER_ID)
        keys = {s.sheet_key for s in contract.sheets}
        assert keys == {F508.SHEET_KEY_F508, F501.SHEET_KEY_F501}

    def test_instrumentation_covers_both_zones(self) -> None:
        """P17 对齐门：instrumentation spec 数 == 受管区数（漏一个就是 D4-35 型事故）。"""
        specs = F5.instrumentation_specs()
        assert len(specs) == len(F5.managed_row_table_specs())
        sheet_keys = {str(s.resolved_sheet_key) for s in specs}
        assert F501.SHEET_KEY_F501 in sheet_keys


class TestF501TbRedLineP20:
    """P20 TB 红线：sync 路径对 `trial_balance` 的写次数必须为 **0**。

    背书：审定表是最容易误写 TB 的一类表（E/I 列就是审定数）。平台铁律是审定数入
    `trial_balance` **只能**走显式发布门 `POST /workpapers/{id}/audit-determination/
    publish-to-tb`（必经二次确认），sync 的 OO↔HTML 往返属"数据变化"，不得触发发布。

    🔴 判定走 AST 而非文本匹配：docstring 里正当地写着「TB 回写走显式发布门」这类说明，
    文本匹配会把说明当违规（判据自己假红）。
    """

    #: sync 路径上的模块（provider + 本 entry 全部受管区 sheet spec）。
    _SYNC_MODULES = (
        "phase5_f5_cost_of_sales.py",
        "phase5_f5_01_adjudication.py",
        "phase5_f5_08_major_adjustment.py",
    )
    #: 禁止出现的写 TB 标识（调用名或字符串字面量）。
    _FORBIDDEN = (
        "publish_to_tb",
        "publishToTb",
        "publish-to-tb",
        "writeback_trial_balance",
        "trial-balance/writeback",
        "trial_balance",
    )

    def _non_docstring_constants_and_calls(self, path) -> tuple[set[str], set[str]]:
        import ast

        tree = ast.parse(path.read_text(encoding="utf-8"))
        docstrings: set[int] = set()
        for node in ast.walk(tree):
            if isinstance(
                node, (ast.Module, ast.FunctionDef, ast.AsyncFunctionDef, ast.ClassDef)
            ):
                body = getattr(node, "body", [])
                if (
                    body
                    and isinstance(body[0], ast.Expr)
                    and isinstance(body[0].value, ast.Constant)
                    and isinstance(body[0].value.value, str)
                ):
                    docstrings.add(id(body[0].value))

        literals: set[str] = set()
        calls: set[str] = set()
        for node in ast.walk(tree):
            if isinstance(node, ast.Constant) and isinstance(node.value, str):
                if id(node) not in docstrings:
                    literals.add(node.value)
            elif isinstance(node, ast.Call):
                func = node.func
                if isinstance(func, ast.Name):
                    calls.add(func.id)
                elif isinstance(func, ast.Attribute):
                    calls.add(func.attr)
        return literals, calls

    def test_sync_modules_never_write_trial_balance(self) -> None:
        base = _BACKEND / "app" / "services" / "workpaper_sync"
        offences: list[str] = []
        for name in self._SYNC_MODULES:
            path = base / name
            assert path.exists(), f"sync 模块缺失: {name}"
            literals, calls = self._non_docstring_constants_and_calls(path)
            for token in self._FORBIDDEN:
                for lit in literals:
                    if token in lit:
                        offences.append(f"{name}: 字符串字面量含 {token!r} → {lit!r}")
                if token in calls:
                    offences.append(f"{name}: 调用了 {token!r}")
        assert offences == [], "sync 路径出现写 TB 痕迹：\n" + "\n".join(offences)

    def test_publish_to_tb_stays_the_only_entrance(self) -> None:
        """前端 `publishToTb` 仍在、且仍指向显式发布门端点（唯一入口未被绕开）。"""
        ts = (
            _BACKEND.parent
            / "audit-platform"
            / "frontend"
            / "src"
            / "components"
            / "workpaper"
            / "composables"
            / "useF5Adjudication.ts"
        )
        assert ts.exists(), "useF5Adjudication.ts 不存在（路径漂移需同步判据）"
        text = ts.read_text(encoding="utf-8")
        assert "audit-determination/publish-to-tb" in text, "显式发布门端点不得消失"
        assert "'6401'" in text or '"6401"' in text, "科目 6401 不得消失"
        assert "occurrence" in text, "发生额口径（amount_kind=occurrence）不得消失"
        # 🔴 必须先剥注释再查旧端点：文件顶部注释正当地记着「原为 trial-balance/writeback
        # 直写，现改为 publish-to-tb」这段变更说明 —— 直接文本匹配会把说明当违规
        # （与本 class docstring 里对 AST 的同一条理由）。
        assert "trial-balance/writeback" not in _strip_ts_comments(text), (
            "旧端点已删，活代码里不得回归"
        )

    def test_publish_is_not_triggered_from_reactive_callbacks(self) -> None:
        """铁律：发布不得在 watch/onMounted/debounce 回调内触发（只能由显式用户动作）。

        判定方式：取 `publishToTb` 函数体之前最近的 `watch(` / `onMounted(` 起始位置，
        确认 `publishToTb` 的定义不在这些回调的词法范围内 —— 这里用保守近似：
        `publishToTb` 必须是顶层 `async function` / `function` 声明。
        """
        import re

        ts = (
            _BACKEND.parent
            / "audit-platform"
            / "frontend"
            / "src"
            / "components"
            / "workpaper"
            / "composables"
            / "useF5Adjudication.ts"
        )
        text = ts.read_text(encoding="utf-8")
        decls = re.findall(r"^(\s*)(?:async\s+)?function\s+publishToTb\b", text, re.M)
        assert decls, "publishToTb 应是具名函数声明（便于静态核查调用点）"
        for indent in decls:
            assert len(indent) <= 2, (
                f"publishToTb 缩进 {len(indent)} 层，疑似嵌在回调内；"
                "发布必须由显式用户动作触发"
            )

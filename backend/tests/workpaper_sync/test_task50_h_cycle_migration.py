# -*- coding: utf-8 -*-
r"""
test_task50_h_cycle_migration — H 循环（除 H1）Excel 独立 entry 迁移验证

spec: workpaper-html-onlyoffice-bidirectional-writeback-closure · Wave 5 Task 50
Requirements: 6.4, 6.5, 12.1, 12.4, 12.10, 12.11, 12.12, 14.1
             （另覆盖 1.4 / 1.6 / 12.8 / 12.9）

验证 Properties:
  - Property 22: 动态列 key 与 label 解耦（真分母：H7 稳定 key 范式的 4 个结构要素正例 +
    全 H 循环 label-as-key 的穷举反例 + 硬编码列数/行数扫描）
  - Property 23: 动态行身份不使用下标（真分母：9 个 primary table 的行身份构造点 +
    17 条位置化命中的三族分类 + 3 条必须**不被**误报的展示序号站点）
  - Property 69: evidence 由逐 scenario 实体与服务端重算闭合（负向分母真验，
    正向逐 scenario 闭合不宣称通过，已登记 BP-4）
  - Property 70: scenario/evidence/contract/bundle 不得跨 entry 复用
  另在真分母上顺带验 Property 28（immutable definition 漂移 fail closed：11 个模板
  digest + 9 个 normalized_structure_hash + 9 组 sheet!cell 真读 + 21 条模板解析）。
  Property 20 / 21 的分母为空 ⇒ **不宣称通过**，见 TestProperty20And21NotClaimed。

═══ 判据设计（避开前四轮踩过的坑，并修正 F/G 对 H 循环不成立的假设）═══

1. **不用空分母重言式**。「H 循环没有 contract 所以 Property 21 通过」是重言式。本文件对
   每条 Property 显式区分「有真分母的部分（真验）」与「无适用分母的部分（只断言前提成立 +
   承载者存在，**不宣称通过**）」，划分与 slice 的 `property_denominators` 逐项对齐。

2. **不用 fail-open**。没有 `pytest.skip`、没有 `if x: assert ...`（缺值即跳过）、
   没有 `except Exception` 吞异常。缺文件/缺字段一律打红。

3. **不用硬编码豁免**。本 slice 没有 pilot，也不给任何 entry 开例外列。

4. **不假设与 F/G 同形**。H 循环实测出六处不同（slice 的 `h_cycle_form_differences`），
   判据逐 entry 从 slice 读声明再回源码核对：
   * 写载体三族（宿主内联 `http` / FormData 用 `api` / 逐 Tab 自持久化）—— F 的
     「持久化 composable 必须自带 GET+PUT」对 6 条 entry 必然失配；
   * 读路径五种，其中 4 条 entry 的载体里**没有** GET /checklist-responses；
   * 主表键命名不统一（H10 无 sheet 尾码 / H5 是前缀拼接 / H7 内联在 Tab），
     身份字段名也分 `rowId` 八条与 `id` 一条；
   * H3/H7 双计量模式；H7 双写路径双客户端；H10 有第三个存储（localStorage 草稿）；
   * 5 条 entry 的 legacy 载体生产零消费或错位消费（宿主内联了第二份实现）。

依赖（缺任一即 fail closed）：
  - backend/data/workpaper_sync_h_cycle_manifest_slice.json（frozen slice）
  - backend/data/workpaper_sync_h_cycle_deletion_plan.json（deletion plan）
  - backend/data/workpaper_sync_migration_paradigm.json（Task 45 范式，只读）
  - backend/data/workpaper_sync_entry_manifest.json（source-backed manifest，只读）
  - backend/wp_templates/H/ 与 backend/wp_templates/_index.json（运行时权威，唯一真源）
"""
from __future__ import annotations

import hashlib
import json
import pathlib
import re
from typing import Any, Callable

import pytest

# ────────────────────────────────────────────────────────────────────────────
# Paths
# ────────────────────────────────────────────────────────────────────────────
# __file__ = backend/tests/workpaper_sync/test_...py
# parents[0] = workpaper_sync/, [1] = tests/, [2] = backend/, [3] = repo root
_THIS = pathlib.Path(__file__).resolve()
ROOT = _THIS.parents[3]
BACKEND = ROOT / "backend"
FRONTEND = ROOT / "audit-platform" / "frontend" / "src"
WP_COMPONENTS = FRONTEND / "components" / "workpaper"
COMPOSABLES = WP_COMPONENTS / "composables"
SYNC_DIR = WP_COMPONENTS / "sync"
DATA = BACKEND / "data"

MANIFEST_SLICE_PATH = DATA / "workpaper_sync_h_cycle_manifest_slice.json"
DELETION_PLAN_PATH = DATA / "workpaper_sync_h_cycle_deletion_plan.json"
PARADIGM_PATH = DATA / "workpaper_sync_migration_paradigm.json"
FULL_MANIFEST_PATH = DATA / "workpaper_sync_entry_manifest.json"
D_SLICE_PATH = DATA / "workpaper_sync_d_cycle_manifest_slice.json"
E_SLICE_PATH = DATA / "workpaper_sync_e_cycle_manifest_slice.json"
F_SLICE_PATH = DATA / "workpaper_sync_f_cycle_manifest_slice.json"
G_SLICE_PATH = DATA / "workpaper_sync_g_cycle_manifest_slice.json"
CONTRACT_DIR = DATA / "workpaper_sync_contracts"
H1_PILOT_CONTRACT = CONTRACT_DIR / "h1.disposal_check.json"
TEMPLATE_DIR = BACKEND / "wp_templates"
TEMPLATE_INDEX = TEMPLATE_DIR / "_index.json"
CHECKLIST_ROUTER = BACKEND / "app" / "routers" / "checklist_responses.py"
REGISTRY = BACKEND / "app" / "services" / "workpaper_sync" / "adapters" / "registry.py"
HTML_RENDERER_REGISTRY = WP_COMPONENTS / "htmlRendererRegistry.ts"
SHARED_BASE = COMPOSABLES / "useWorkpaperEntryDualMode.ts"
H7_LISTED_MODEL = COMPOSABLES / "h7ListedDisclosureModel.ts"
H7_SOE_MODEL = COMPOSABLES / "h7SoeDisclosureModel.ts"
H8_SYNC_PAYLOAD = COMPOSABLES / "h8DisclosureSyncPayload.ts"

#: AC 1.3 的能力态枚举（真源在范式 JSON 的 adjudication_criteria.capability_enum）。
CAPABILITY_ENUM = ("bidirectional", "single_html", "single_onlyoffice", "unreachable")
#: step 3 的二值结论（真源在范式 JSON 的 anti_patterns[AP-1].allowed_verdict_values）。
HTML_COUNTERPART_VERDICTS = ("none", "exists")

#: AC 1.4 的 UI 义务落点（Task 46 收口新建的单一真源，本任务复用不新建第二份）。
NOTICE_MODULE = SYNC_DIR / "workpaperEntrySyncNotice.ts"
NOTICE_COMPONENT = SYNC_DIR / "GtEntrySyncCapabilityNotice.vue"
NOTICE_COMPONENT_NAME = "GtEntrySyncCapabilityNotice"

#: H1 的 entry —— 任务正文明令排除（Task 42 的 pilot），本文件多处前提依赖「它不在 slice 里」。
H1_ENTRY_ID = "xlsx/gt-h1-fixed-assets"
#: H1 的宿主，位置化扫描与 label-as-key 扫描都必须排除它（pilot 不在本 slice 分母内）。
H1_HOST_NAME = "GtH1FixedAssets.vue"

#: 本 slice 覆盖的循环子目录（h2..h10），用于组织「全 H 循环」扫描的分母。
H_SUBDIRS = ("h2", "h3", "h4", "h5", "h6", "h7", "h8", "h9", "h10")
#: composables 里属本 slice 的文件名前缀正则（h2..h10，排除 h1/h10 混淆）。
H_COMPOSABLE_RE = re.compile(r"(?:use)?[hH](?:2|3|4|5|6|7|8|9|10)(?![0-9])")


# ────────────────────────────────────────────────────────────────────────────
# Helpers
# ────────────────────────────────────────────────────────────────────────────
def _load(path: pathlib.Path) -> dict:
    return json.loads(path.read_text(encoding="utf-8"))


def _sha256_of(path: pathlib.Path) -> str:
    return hashlib.sha256(path.read_bytes()).hexdigest()


def _strip_ts_comments(source: str) -> str:
    """剥掉 TS/Vue 注释 —— 说明文字不得充当判据证据。

    🔴 剥完必须反向自检（见 TestGuardSelfChecks）。本文件对它的依赖很重：
    「H 循环只有 2 处 label-as-key」这条穷举判据的分母就是剥注释后的正文 ——
    剥不干净会假红（h7ListedDisclosureModel 的注释里逐字写着 `key: label`），
    剥过头会假绿。
    """
    source = re.sub(r"/\*[\s\S]*?\*/", "", source)
    source = re.sub(r"(?m)^\s*//.*$", "", source)
    source = re.sub(r"(?m)//[^\n\"'`]*$", "", source)
    source = re.sub(r"<!--[\s\S]*?-->", "", source)
    return source


def _vue_template(source: str) -> str:
    """取 SFC 的**外层** template 区块。

    🔴 不能用 `source.split("</template>")[0]`：SFC 里嵌套的具名插槽 / `<template v-else>`
    会先闭合，第一处 `</template>` 落在内层 ⇒ 截出来的片段比真实模板短，判据会把
    「已在模板里」误判成「不在模板里」。这里改用 `<script` 边界。
    """
    marker = source.find("<script")
    assert marker > 0, "找不到 <script 边界 —— 该文件不是常规 SFC（template 在 script 之前）"
    head = source[:marker]
    assert "<template>" in head, "SFC 头部没有 <template>"
    return head


def _toolbar_block(template: str, gate_anchor: str) -> str:
    """按 slice 声明的工具栏锚点截出该区块（开标签行 → 同缩进的 `</div>`）。

    🔴 为什么按 slice 声明取锚点而不是写死正则：H 循环 9 个宿主的工具栏 class 与门控
    表达式逐个不同（h2-header-toolbar 带 `v-if="showModeSwitch"`、h5 带
    `showHtmlToolbar && currentSheet !== 'H5'` 复合门控、h10 是 `h10-toolbar` 且按中文
    sheet 名排除）。写死任一种都会让其余的判据静默恒真（挂载点根本不在被搜的区块里也不报错）。
    """
    lines = template.splitlines()
    opens = [i for i, ln in enumerate(lines) if gate_anchor in ln]
    assert len(opens) == 1, (
        f"工具栏锚点 {gate_anchor!r} 在模板里命中 {len(opens)} 次（应为 1）"
    )
    start = opens[0]
    indent = len(lines[start]) - len(lines[start].lstrip())
    for i in range(start + 1, len(lines)):
        if lines[i].strip() == "</div>" and (len(lines[i]) - len(lines[i].lstrip())) == indent:
            return "\n".join(lines[start : i + 1])
    pytest.fail(f"锚点 {gate_anchor!r} 之后找不到同缩进的 </div> —— 工具栏区块无法界定")


def _cell_value(path: pathlib.Path, sheet: str, cell: str) -> Any:
    """openpyxl 真读权威模板的一个单元格（data_only）。"""
    import openpyxl
    from openpyxl.utils import column_index_from_string

    workbook = openpyxl.load_workbook(path, data_only=True)
    try:
        assert sheet in workbook.sheetnames, (
            f"{path.name} 里没有 sheet {sheet!r}；实有 {workbook.sheetnames}"
        )
        match = re.fullmatch(r"([A-Z]+)(\d+)", cell)
        assert match, f"非法单元格坐标 {cell!r}"
        return workbook[sheet].cell(
            row=int(match.group(2)),
            column=column_index_from_string(match.group(1)),
        ).value
    finally:
        workbook.close()


def _sheet_names(path: pathlib.Path) -> list[str]:
    import openpyxl

    workbook = openpyxl.load_workbook(path, read_only=True)
    try:
        return list(workbook.sheetnames)
    finally:
        workbook.close()


def _indexed_relpaths() -> set[str]:
    """`_index.json` 登记的相对路径集合（分隔符归一为 /）。"""
    index = _load(TEMPLATE_INDEX)
    return {str(f["relative_path"]).replace("\\", "/") for f in index["files"]}


def _resolve_repo(rel: str) -> pathlib.Path:
    """把 slice 里登记的仓库相对路径（可带 `#Lxx` 或 `!sheet!cell`）解析成绝对路径。"""
    head = rel.split("#")[0].split("!")[0]
    return ROOT / head


def _value_expr_after_key(line: str, key: str) -> str | None:
    """截出 `key:` 之后**属于它自己**的值表达式（到同层逗号 / 行尾止）。

    🔴 为什么不能拿整行当判据：H3/H6 的载入行是
    `{ ...raw, rowId: raw.rowId ?? gen(), seq: raw.seq ?? idx + 1 }` —— 整行含
    `idx + 1`，但那是 `seq`（展示序号，每次载入重算）的值，不是 `rowId` 的。按整行判会把
    三条 entry 误判成位置化行身份（G 循环在 G12 上踩过同一个坑）。
    """
    match = re.search(r"\b" + re.escape(key) + r"\s*\??\s*:", line)
    if not match:
        return None
    rest = line[match.end() :]
    depth = 0
    in_tmpl = False
    for pos, char in enumerate(rest):
        if char == "`":
            in_tmpl = not in_tmpl
            continue
        if in_tmpl:
            continue
        if char in "([{":
            depth += 1
        elif char in ")]}":
            if depth == 0:
                return rest[:pos]
            depth -= 1
        elif char == "," and depth == 0:
            return rest[:pos]
    return rest


#: 位置化身份 token —— 命中即说明身份值由数组下标 / 序号派生（forbidden 三种之一）。
_POSITIONAL_IDENTITY_TOKEN = re.compile(
    r"\$\{\s*i\s*\}"
    r"|\$\{\s*idx\s*\}"
    r"|\$\{\s*index\s*\}"
    r"|\$\{\s*seq\s*\}"
    r"|(?:\|\||\?\?)\s*(?:i|idx|index)\s*\+\s*1"
    r"|(?:\|\||\?\?)\s*String\(\s*(?:i|idx|index)"
    r"|\bindexOf\("
)

#: 身份字段名候选 —— H 循环实测两族（rowId 八条 / id 一条），另留 rowKey 供将来。
_IDENTITY_KEYS = ("rowId", "rowKey", "id")

#: 「列 key 用 label」两种形态（Property 22 的反例分母）。
_COLUMN_KEY_IS_LABEL = re.compile(r"\bkey\s*:\s*[A-Za-z_$][\w$]*(?:\.[\w$]+)*\.label\b")
_ROW_CELL_KEY_IS_LABEL = re.compile(r"\brow\s*\[\s*[A-Za-z_$][\w$]*(?:\.[\w$]+)*\.label\s*\]")
#: 「动态区骨架行数写死」形态：`blankRows(<expr>, <整数字面量>)`。
_BLANK_ROWS_FIXED = re.compile(r"\bblankRows\s*\(\s*[^,()]+,\s*\d+\s*\)")
#: 「按公司/单位横向展开写死列名」形态。
_HORIZONTAL_COLUMN_LITERAL = re.compile(r"['\"`](?:公司|单位)[1-9]\d*['\"`]")


def _h_cycle_files() -> list[pathlib.Path]:
    """「全 H 循环」扫描的分母：h2..h10 子目录 + 本 slice 的 composables + GtH*.vue（排除 H1 宿主）。

    分母**现扫**而不是写死清单 —— 新增文件自动进入，改判据不必改守卫。
    """
    out: set[pathlib.Path] = set()
    for name in H_SUBDIRS:
        directory = WP_COMPONENTS / name
        assert directory.is_dir(), f"缺 H 循环子目录 {directory}"
        out |= {p for p in directory.rglob("*") if p.is_file() and p.suffix in (".ts", ".vue")}
    for path in COMPOSABLES.rglob("*"):
        if path.is_file() and path.suffix == ".ts" and H_COMPOSABLE_RE.match(path.name):
            out.add(path)
    out |= {p for p in WP_COMPONENTS.glob("GtH*.vue") if p.name != H1_HOST_NAME}
    return sorted(out)


def _positional_identity_hits(files: list[pathlib.Path]) -> list[tuple[str, int, str, str]]:
    """穷举「身份字段的值由位置派生」的命中，返回 `(相对路径, 行号, 身份键, 表达式)`。"""
    hits: list[tuple[str, int, str, str]] = []
    for path in files:
        source = path.read_text(encoding="utf-8", errors="replace")
        rel = path.relative_to(ROOT).as_posix()
        for no, line in enumerate(source.splitlines(), 1):
            for key in _IDENTITY_KEYS:
                expr = _value_expr_after_key(line, key)
                if expr is not None and _POSITIONAL_IDENTITY_TOKEN.search(expr):
                    hits.append((rel, no, key, expr.strip()))
                    break
    return hits


def _label_key_hits(files: list[pathlib.Path]) -> dict[str, list[str]]:
    """穷举「label 当 key」的命中（剥注释后），按两种形态分别返回 `文件#Lxx` 列表。"""
    out: dict[str, list[str]] = {"column_key_is_label": [], "row_cell_key_is_label": []}
    for path in files:
        body = _strip_ts_comments(path.read_text(encoding="utf-8", errors="replace"))
        rel = path.relative_to(ROOT).as_posix()
        for no, line in enumerate(body.splitlines(), 1):
            if _COLUMN_KEY_IS_LABEL.search(line):
                out["column_key_is_label"].append(f"{rel}#L{no}")
            if _ROW_CELL_KEY_IS_LABEL.search(line):
                out["row_cell_key_is_label"].append(f"{rel}#L{no}")
    return out


def _import_specifier_consumers(module_stem: str) -> list[str]:
    """全前端 src 里 import 路径字面量含 `module_stem` 的文件（排除自身），用于孤儿判定。"""
    pattern = re.compile(r"['\"][^'\"]*" + re.escape(module_stem) + r"['\"]")
    out: list[str] = []
    for path in FRONTEND.rglob("*"):
        if not (path.is_file() and path.suffix in (".ts", ".vue")):
            continue
        if path.stem == module_stem:
            continue
        if pattern.search(path.read_text(encoding="utf-8", errors="replace")):
            out.append(path.relative_to(ROOT).as_posix())
    return sorted(out)


def _strip_null_placeholder_columns(window: str) -> str:
    """把 `conclusion: null` / `remark: undefined` 这类**显式空占位**从写入点窗口剔除。

    🔴 `payload_column_mode` 判的是「业务内容写进哪一列」，而 `conclusion: null` 恰恰是
    「**不**往这列写内容」。H7 的主表写入点是
    `items: [{ item_id: itemId, conclusion: null, remark }]` —— 只看 token `conclusion:`
    在不在窗口里，会把它误判成 dual_write（G 循环 G2 踩过同一个坑）。剔除只针对**字面
    null / undefined**，真双写（`conclusion: existing?.conclusion ?? null`、H9 的 ES 简写
    `conclusion`）一个都不会被剔掉。
    """
    return re.sub(r"\b(remark|conclusion)\s*:\s*(?:null|undefined)\s*,?", "", window)


def _write_site_window(path: pathlib.Path, line_no: int, span: int = 4) -> str:
    lines = path.read_text(encoding="utf-8", errors="replace").splitlines()
    lo = max(0, line_no - 1 - span)
    hi = min(len(lines), line_no + span)
    return "\n".join(lines[lo:hi])


def _line_no_of(ref: str) -> int:
    match = re.search(r"#L(\d+)", ref)
    assert match, f"引用 {ref!r} 里没有 #Lxx 行号"
    return int(match.group(1))


def _client_probe(client: str, verb: str) -> Callable[[str], bool]:
    """按 slice 登记的 HTTP 客户端标识拼探针。

    🔴 不写成「`api.` 或 `http.` 二选一」的固定并集：那样任一 entry 的绑定名改了都不会红
    （另一族还在，并集恒真）。这里按**该 entry 自己声明的**标识拼正则 —— 声明与源码不一致
    立刻打红。
    """
    esc = re.escape(client)

    def probe(source: str) -> bool:
        return bool(re.search(rf"{esc}\.{verb}\(\s*`[^`]*checklist-responses`", source))

    return probe


# ────────────────────────────────────────────────────────────────────────────
# Fixtures
# ────────────────────────────────────────────────────────────────────────────
@pytest.fixture(scope="module")
def manifest_slice() -> dict:
    return _load(MANIFEST_SLICE_PATH)


@pytest.fixture(scope="module")
def deletion_plan() -> dict:
    return _load(DELETION_PLAN_PATH)


@pytest.fixture(scope="module")
def paradigm() -> dict:
    return _load(PARADIGM_PATH)


@pytest.fixture(scope="module")
def full_manifest() -> dict:
    return _load(FULL_MANIFEST_PATH)


@pytest.fixture(scope="module")
def h_files() -> list[pathlib.Path]:
    return _h_cycle_files()


# ════════════════════════════════════════════════════════════════════════════
# 判据自检 —— 判据本身不能恒真
# ════════════════════════════════════════════════════════════════════════════
class TestGuardSelfChecks:
    """反向自检：辅助函数既不能什么都不做，也不能把代码一起剥掉/截掉。"""

    def test_strip_comments_removes_comments_but_keeps_code(self) -> None:
        sample = (
            "// leading line comment\n"
            "/* block */\n"
            "const KEEP = 'value' // trailing\n"
            "<!-- html comment -->\n"
        )
        out = _strip_ts_comments(sample)
        assert "leading line comment" not in out
        assert "block" not in out
        assert "html comment" not in out
        assert "const KEEP = 'value'" in out.replace("\n", " ").strip() or "KEEP" in out

    def test_strip_comments_does_not_eat_string_literals(self) -> None:
        sample = "const URL = 'https://example.com/a//b'\n"
        assert "https://example.com/a//b" in _strip_ts_comments(sample)

    def test_strip_comments_really_hides_the_h7_paradigm_comment(self) -> None:
        """H7 的注释里逐字写着 `key: label` —— 剥不干净会让 label-as-key 穷举假红。"""
        raw = H7_LISTED_MODEL.read_text(encoding="utf-8")
        assert "key: label" in raw, "H7 注释形态变了，本自检的前提失效，请重写"
        assert not _COLUMN_KEY_IS_LABEL.search(_strip_ts_comments(raw)), (
            "剥注释后 H7 模型仍被 label-as-key 判据命中 ⇒ 剥得不干净"
        )

    def test_vue_template_extraction_does_not_stop_at_inner_template(self) -> None:
        sample = (
            "<template>\n"
            "  <div>\n"
            "    <template #default><span>x</span></template>\n"
            "    <MARKER />\n"
            "  </div>\n"
            "</template>\n"
            "<script setup lang=\"ts\">\n"
            "const a = 1\n"
            "</script>\n"
        )
        assert "<MARKER />" in _vue_template(sample)

    def test_toolbar_block_extraction_is_bounded(self) -> None:
        template = (
            "<template>\n"
            "  <div>\n"
            "      <div class=\"x-toolbar\">\n"
            "        <A />\n"
            "      </div>\n"
            "      <div class=\"other\">\n"
            "        <B />\n"
            "      </div>\n"
            "  </div>\n"
            "</template>\n"
        )
        block = _toolbar_block(template, 'class="x-toolbar"')
        assert "<A />" in block
        assert "<B />" not in block

    def test_toolbar_block_rejects_ambiguous_anchor(self) -> None:
        template = (
            "<template>\n"
            "      <div class=\"dup\">\n"
            "      </div>\n"
            "      <div class=\"dup\">\n"
            "      </div>\n"
            "</template>\n"
        )
        with pytest.raises(AssertionError):
            _toolbar_block(template, 'class="dup"')

    def test_value_expr_after_key_is_scoped_to_its_own_key(self) -> None:
        line = "return { ...raw, rowId: raw.rowId ?? gen(), seq: raw.seq ?? idx + 1 }"
        row_expr = _value_expr_after_key(line, "rowId")
        seq_expr = _value_expr_after_key(line, "seq")
        assert row_expr is not None and seq_expr is not None
        assert not _POSITIONAL_IDENTITY_TOKEN.search(row_expr), (
            "rowId 的表达式被误判为位置化 ⇒ 判据回到了「整行判」的错口径"
        )
        assert _POSITIONAL_IDENTITY_TOKEN.search(seq_expr), (
            "seq 的表达式没被识别出位置化 ⇒ 判据整体失灵（连真的都抓不到）"
        )

    def test_positional_token_catches_the_real_seed_form(self) -> None:
        expr = _value_expr_after_key("rowId: `seed-${idx}`,", "rowId")
        assert expr is not None and _POSITIONAL_IDENTITY_TOKEN.search(expr)

    def test_null_placeholder_stripper_only_removes_literal_nulls(self) -> None:
        real_dual = "{ item_id: id, conclusion: existing?.conclusion ?? null, remark: v }"
        placeholder = "{ item_id: id, conclusion: null, remark: v }"
        assert "conclusion" in _strip_null_placeholder_columns(real_dual)
        assert "conclusion" not in _strip_null_placeholder_columns(placeholder)
        assert "remark" in _strip_null_placeholder_columns(placeholder)

    def test_client_probe_is_client_specific(self) -> None:
        src = "await api.put(`/api/workpapers/${id}/checklist-responses`, {})"
        assert _client_probe("api", "put")(src)
        assert not _client_probe("http", "put")(src), (
            "探针对另一族客户端也命中 ⇒ 声明与源码不一致时不会打红"
        )

    def test_h_cycle_denominator_is_non_vacuous_and_excludes_the_pilot(
        self, h_files: list[pathlib.Path]
    ) -> None:
        assert len(h_files) >= 300, f"H 循环扫描分母只有 {len(h_files)} 个文件 —— 疑似 glob 失效"
        names = {p.name for p in h_files}
        assert H1_HOST_NAME not in names, "H1 宿主进了分母 —— pilot 必须排除"
        assert "GtH8RightOfUseAssets.vue" in names
        assert "h7ListedDisclosureModel.ts" in names

    def test_all_required_artifacts_exist(self) -> None:
        for path in (
            MANIFEST_SLICE_PATH,
            DELETION_PLAN_PATH,
            PARADIGM_PATH,
            FULL_MANIFEST_PATH,
            D_SLICE_PATH,
            E_SLICE_PATH,
            F_SLICE_PATH,
            G_SLICE_PATH,
            TEMPLATE_INDEX,
            CHECKLIST_ROUTER,
            REGISTRY,
            HTML_RENDERER_REGISTRY,
            NOTICE_MODULE,
            NOTICE_COMPONENT,
            SHARED_BASE,
            H7_LISTED_MODEL,
            H7_SOE_MODEL,
            H8_SYNC_PAYLOAD,
            H1_PILOT_CONTRACT,
        ):
            assert path.exists(), f"缺依赖 {path}"


# ════════════════════════════════════════════════════════════════════════════
# 裁决合法性（AC 12.8 / 12.9 / 1.3 / 12.1 + AP-1）
# ════════════════════════════════════════════════════════════════════════════
class TestAdjudicationLegality:
    """**Validates: Requirements 12.1, 12.8, 12.9**

    裁 single_onlyoffice 的唯一合法判据是「无 HTML 对端」；理由不得是循环论证；
    capability 要么落在枚举内，要么是待裁决态且三字段齐备。
    """

    def test_every_entry_has_a_binary_html_counterpart_verdict(
        self, manifest_slice: dict
    ) -> None:
        for entry in manifest_slice["independent_entries"]:
            verdict = entry.get("html_counterpart_verdict")
            assert verdict in HTML_COUNTERPART_VERDICTS, (
                f"{entry['entry_id']}: html_counterpart_verdict={verdict!r} 不是二值结论"
                "（unresolved / unknown / 空值是「还没查」，见 AP-3）"
            )
            refs = entry.get("html_counterpart_source_refs")
            assert isinstance(refs, list) and refs, f"{entry['entry_id']}: 结论缺 source_refs"

    def test_single_onlyoffice_requires_no_html_counterpart(self, manifest_slice: dict) -> None:
        for entry in manifest_slice["independent_entries"]:
            if entry.get("capability") == "single_onlyoffice":
                assert entry["html_counterpart_verdict"] == "none", (
                    f"{entry['entry_id']}: 有 HTML 对端却裁 single_onlyoffice（违反 AC 12.8）"
                )

    def test_adjudication_reason_is_not_circular(self, manifest_slice: dict, paradigm: dict) -> None:
        """AP-1：理由不得**只**由「本任务应交付的产物尚不存在」构成。"""
        markers = [
            str(m).lower()
            for ap in paradigm["adjudication_criteria"]["anti_patterns"]
            if ap.get("id") == "AP-1"
            for m in (ap.get("circular_reason_markers") or [])
        ]
        assert markers, "范式里取不到 AP-1 的 circular_reason_markers —— 判据无分母"
        for entry in manifest_slice["independent_entries"]:
            reason = str(entry["adjudication"]["reason"]).lower()
            hit = [m for m in markers if m in reason]
            if not hit:
                continue
            assert "html 对端" in reason or "html_counterpart_verdict" in reason, (
                f"{entry['entry_id']}: 理由命中循环论证标记 {hit[:3]} 且通篇没有对端结论"
            )

    def test_capability_is_enum_or_explicitly_pending(self, manifest_slice: dict) -> None:
        pending = 0
        for entry in manifest_slice["independent_entries"]:
            cap = entry.get("capability")
            if cap is None:
                pending += 1
                assert isinstance(entry.get("capability_verdict_stage"), str) and entry[
                    "capability_verdict_stage"
                ].strip(), f"{entry['entry_id']}: 待裁决态缺非空 capability_verdict_stage"
                assert entry.get("capability_target") in CAPABILITY_ENUM, (
                    f"{entry['entry_id']}: capability_target 必须落在 capability_enum 内"
                )
                blocked = entry.get("capability_target_blocked_by")
                assert isinstance(blocked, list) and blocked, (
                    f"{entry['entry_id']}: capability_target_blocked_by 必须是非空数组"
                )
            else:
                assert cap in CAPABILITY_ENUM, f"{entry['entry_id']}: capability={cap!r} 不在枚举内"
        counters = manifest_slice["honest_adjudication_summary"]["slice_counters"]
        assert counters["unadjudicated"] == pending, (
            f"SR-9: unadjudicated={counters['unadjudicated']} != 待裁决实际条数 {pending}"
        )

    def test_capability_matches_honest_capability(self, manifest_slice: dict) -> None:
        for entry in manifest_slice["independent_entries"]:
            assert entry["capability"] == entry["adjudication"]["honest_capability"], (
                f"{entry['entry_id']}: 对外 capability 与 adjudication.honest_capability 双口径"
            )

    def test_single_or_pending_entries_carry_no_identity(self, manifest_slice: dict) -> None:
        identity = (
            "adapter_id",
            "authority_model",
            "definition_bundle",
            "instrumentation_candidate",
            "published_representation",
        )
        for entry in manifest_slice["independent_entries"]:
            cap = entry["capability"]
            if cap == "bidirectional":
                continue
            non_null = [k for k in identity if entry.get(k) is not None]
            assert not non_null, (
                f"{entry['entry_id']}: capability={cap!r} 却挂着 {non_null}（AP-5 / SR-5）"
            )

    def test_adjudication_carries_both_negative_reasons(self, manifest_slice: dict) -> None:
        for entry in manifest_slice["independent_entries"]:
            adj = entry["adjudication"]
            for key in ("not_single_html_because", "not_bidirectional_because"):
                assert isinstance(adj.get(key), str) and len(adj[key]) > 40, (
                    f"{entry['entry_id']}: {key} 必须写明理由，不能是空话"
                )

    def test_capability_blockers_reference_real_preconditions(self, manifest_slice: dict) -> None:
        ids = {bp["id"] for bp in manifest_slice["blocking_preconditions"]}
        for entry in manifest_slice["independent_entries"]:
            for blocker in entry["capability_target_blocked_by"]:
                assert blocker in ids, (
                    f"{entry['entry_id']}: blocked_by 里的 {blocker!r} 不在 blocking_preconditions"
                )

    def test_every_entry_scoped_blocker_has_a_carrier(self, manifest_slice: dict) -> None:
        """反向：声明为「阻断某 entry 的 bidirectional」的 BP 必须真被某个 entry 引用。

        BP-9（manifest 一致性）与 BP-10（UI 义务）blocks 的不是 bidirectional，故按
        `blocks` 字段是否含 'bidirectional' 现算分母，而不是写死排除列表。
        """
        referenced = {
            b
            for entry in manifest_slice["independent_entries"]
            for b in entry["capability_target_blocked_by"]
        }
        for bp in manifest_slice["blocking_preconditions"]:
            if "bidirectional" not in str(bp["blocks"]):
                continue
            assert bp["id"] in referenced, (
                f"{bp['id']} 声称阻断 bidirectional，却没有任何 entry 把它列进 blocked_by"
                " ⇒ 哑登记"
            )

    def test_manifest_mirror_divergence_is_registered_not_silently_equal(
        self, manifest_slice: dict, full_manifest: dict
    ) -> None:
        """source manifest 的 capability/html_store 是 overlay 组件级默认值，不是裁决真源。

        🔴 判据是「必须不一致且该不一致已登记」，而不是「必须相等」—— 后者会把 overlay
        默认值当成裁决结论。
        """
        by_id = {e["entry_id"]: e for e in full_manifest["entries"]}
        for entry in manifest_slice["independent_entries"]:
            src = by_id[entry["entry_id"]]
            mirror = entry["manifest_mirror"]
            assert mirror["capability"] == src["capability"], (
                f"{entry['entry_id']}: manifest_mirror.capability 与 source manifest 不符"
            )
            assert mirror["html_store"] == src["html_store"], (
                f"{entry['entry_id']}: manifest_mirror.html_store 与 source manifest 不符"
            )
            assert mirror["capability"] != entry["capability"], (
                f"{entry['entry_id']}: mirror 与 slice 的 capability 相等 ⇒ BP-9 的前提消失"
            )
            assert "BP-9" in str(mirror["divergence_from_slice"]), (
                f"{entry['entry_id']}: 不一致未指向 BP-9"
            )
            assert mirror["legacy_reasons"] == src["evidence"]["legacy_reasons"]


# ════════════════════════════════════════════════════════════════════════════
# HTML 对端结论必须 source-backed（step 3）
# ════════════════════════════════════════════════════════════════════════════
class TestHtmlCounterpartIsSourceBacked:
    """**Validates: Requirements 12.8, 12.11, 14.1**

    每条 entry 的对端结论逐项回源码核对：载体归属 / 客户端 / 读写端点 / payload 列 /
    传输键 / owner 常量 / 模板单元格。**逐 entry 按 slice 声明取判据**，不假设同形。
    """

    def test_source_refs_point_at_real_paths(self, manifest_slice: dict) -> None:
        for entry in manifest_slice["independent_entries"]:
            for ref in entry["html_counterpart_source_refs"]:
                assert _resolve_repo(ref).exists(), f"{entry['entry_id']}: source_ref 不存在 {ref}"

    def test_html_store_endpoints_exist_in_the_router(self, manifest_slice: dict) -> None:
        router = CHECKLIST_ROUTER.read_text(encoding="utf-8")
        lines = router.splitlines()
        for entry in manifest_slice["independent_entries"]:
            source = entry["html_counterpart"]["endpoint_source"]
            for ref in re.findall(r"#L(\d+)", source):
                idx = int(ref) - 1
                assert 0 <= idx < len(lines), f"{entry['entry_id']}: 端点行号越界 {source}"
                assert "@router." in lines[idx], (
                    f"{entry['entry_id']}: endpoint_source 指向的 L{ref} 不是路由装饰器"
                    f"（实际 {lines[idx].strip()!r}）"
                )

    def test_write_carrier_client_and_put_site_agree_with_the_source(
        self, manifest_slice: dict
    ) -> None:
        for entry in manifest_slice["independent_entries"]:
            hc = entry["html_counterpart"]
            carrier = _resolve_repo(hc["write_carrier_path"])
            assert carrier.exists(), f"{entry['entry_id']}: write_carrier_path 不存在"
            source = carrier.read_text(encoding="utf-8")
            client = hc["write_client"]
            assert hc["write_client_import"] in source, (
                f"{entry['entry_id']}: 载体里找不到声明的 client import {hc['write_client_import']!r}"
            )
            assert _client_probe(client, "put")(source), (
                f"{entry['entry_id']}: 载体 {hc['write_carrier_path']} 里没有 "
                f"{client}.put(`...checklist-responses`)"
            )
            put_line = _line_no_of(hc["endpoint_write_source"])
            lines = source.splitlines()
            assert f"{client}.put(" in lines[put_line - 1], (
                f"{entry['entry_id']}: endpoint_write_source 指向的 L{put_line} 不是 "
                f"{client}.put（实际 {lines[put_line - 1].strip()[:90]!r}）"
            )

    def test_read_carrier_matches_its_declared_family(self, manifest_slice: dict) -> None:
        """五种读路径各有自己的判据 —— 写死 `/checklist-responses` 双端点会对 4 条 entry 假红。"""
        for entry in manifest_slice["independent_entries"]:
            hc = entry["html_counterpart"]
            kind = hc["read_carrier"]
            path = _resolve_repo(hc["read_carrier_path"])
            assert path.exists(), f"{entry['entry_id']}: read_carrier_path 不存在"
            source = path.read_text(encoding="utf-8")
            client = hc["write_client"]
            if kind in ("formdata_composable", "per_tab_formdata_instance", "per_tab_self_persisting"):
                assert _client_probe(client, "get")(source), (
                    f"{entry['entry_id']}: 声明 {kind} 却找不到 {client}.get(`...checklist-responses`)"
                )
            elif kind == "render_config_snapshot_passthrough":
                assert "props.htmlData" in source, (
                    f"{entry['entry_id']}: 声明走 render-config 快照，源码里却没有 props.htmlData"
                )
                assert not _client_probe(client, "get")(source), (
                    f"{entry['entry_id']}: 声明「宿主自身不 GET /checklist-responses」，"
                    "但源码里有 —— 声明与实况脱节"
                )
            elif kind == "host_inline_render_config_refetch":
                assert re.search(r"\.get\(\s*`[^`]*render-config`", source), (
                    f"{entry['entry_id']}: 声明 render-config 强制刷新，源码里没有该 GET"
                )
                assert "force_component_type" in source, (
                    f"{entry['entry_id']}: render-config 刷新缺 force_component_type"
                )
                assert not _client_probe(client, "get")(source), (
                    f"{entry['entry_id']}: 声明读路径不是 /checklist-responses，但源码里有该 GET"
                )
            else:
                pytest.fail(f"{entry['entry_id']}: 未知 read_carrier={kind!r}，判据 fail closed")

    def test_read_carrier_families_cover_every_entry_exactly_once(
        self, manifest_slice: dict
    ) -> None:
        """HD-2 的分族登记必须与逐 entry 声明等值（防摘要与明细脱节）。"""
        declared = manifest_slice["h_cycle_form_differences"]["differences"]
        hd2 = next(d for d in declared if d["id"] == "HD-2")
        buckets = hd2["entries_by_read_carrier"]
        flat: list[str] = [e for v in buckets.values() for e in v]
        assert len(flat) == len(set(flat)), "HD-2 的分族有重复 entry"
        actual: dict[str, list[str]] = {}
        for entry in manifest_slice["independent_entries"]:
            actual.setdefault(entry["html_counterpart"]["read_carrier"], []).append(
                entry["entry_id"]
            )
        assert {k: sorted(v) for k, v in buckets.items()} == {
            k: sorted(v) for k, v in actual.items()
        }, "HD-2 的 entries_by_read_carrier 与逐 entry 声明不一致"

    def test_write_carrier_families_cover_every_entry_exactly_once(
        self, manifest_slice: dict
    ) -> None:
        declared = manifest_slice["h_cycle_form_differences"]["differences"]
        hd1 = next(d for d in declared if d["id"] == "HD-1")
        buckets = hd1["entries_by_family"]
        flat: list[str] = [e for v in buckets.values() for e in v]
        assert len(flat) == len(set(flat)) == len(manifest_slice["independent_entries"]), (
            "HD-1 的分族与 entry 总数不符"
        )
        actual: dict[str, list[str]] = {}
        for entry in manifest_slice["independent_entries"]:
            actual.setdefault(entry["html_counterpart"]["write_carrier"], []).append(
                entry["entry_id"]
            )
        assert {k: sorted(v) for k, v in buckets.items()} == {
            k: sorted(v) for k, v in actual.items()
        }, "HD-1 的 entries_by_family 与逐 entry 声明不一致"

    def test_payload_column_mode_matches_the_write_site(self, manifest_slice: dict) -> None:
        for entry in manifest_slice["independent_entries"]:
            hc = entry["html_counterpart"]
            path = _resolve_repo(hc["payload_write_site"])
            window = _write_site_window(path, _line_no_of(hc["payload_write_site"]))
            stripped = _strip_null_placeholder_columns(window)
            has_remark = "remark" in stripped
            has_conclusion = "conclusion" in stripped
            mode = hc["payload_column_mode"]
            if mode == "dual_write_remark_and_conclusion":
                assert has_remark and has_conclusion, (
                    f"{entry['entry_id']}: 声明双写，剔除空占位后窗口里 "
                    f"remark={has_remark} conclusion={has_conclusion}"
                )
            elif mode == "remark_only_with_explicit_null_conclusion":
                assert has_remark and not has_conclusion, (
                    f"{entry['entry_id']}: 声明「remark + 显式 null 占位」，实际 "
                    f"remark={has_remark} conclusion={has_conclusion}"
                )
                assert "conclusion" in window, (
                    f"{entry['entry_id']}: 声明有 conclusion 空占位，原窗口里却没有该 token"
                )
                assert hc.get("payload_null_placeholder_columns") == ["conclusion"], (
                    f"{entry['entry_id']}: 空占位列未双向登记"
                )
            else:
                pytest.fail(f"{entry['entry_id']}: 未知 payload_column_mode={mode!r}")
            assert hc["payload_column"] == "remark", (
                f"{entry['entry_id']}: H 循环 9 条 entry 的业务载荷列实测全是 remark"
            )

    def test_transport_key_kind_matches_the_source(self, manifest_slice: dict) -> None:
        """HD-4：主表键的**声明形态**必须能在 owner 里找到对应写法。"""
        for entry in manifest_slice["independent_entries"]:
            table = entry["html_counterpart"]["primary_table"]
            owner = _resolve_repo(table["owner_module"])
            assert owner.exists(), f"{entry['entry_id']}: owner_module 不存在 {table['owner_module']}"
            source = owner.read_text(encoding="utf-8")
            kind = table["owner_declaration_kind"]
            item_id = table["item_id"]
            if kind == "module_constant":
                const = table["owner_constant"]
                assert re.search(
                    rf"\b(?:const|let)\s+{re.escape(const)}\s*=\s*['\"]{re.escape(item_id)}['\"]",
                    source,
                ), f"{entry['entry_id']}: 找不到 `const {const} = '{item_id}'`"
            elif kind == "module_constant_prefix":
                const = table["owner_constant"]
                expr = table["owner_composition_expression"]
                assert re.search(rf"\b(?:const|let)\s+{re.escape(const)}\s*=\s*['\"]", source), (
                    f"{entry['entry_id']}: 找不到前缀常量 {const}"
                )
                assert expr in source, (
                    f"{entry['entry_id']}: 找不到拼接式 {expr}（按字面量 grep 会假红）"
                )
            elif kind == "inline_literal_in_tab_component":
                assert table["owner_constant"] is None, (
                    f"{entry['entry_id']}: 内联字面量形态不应声明 owner_constant"
                )
                assert f"'{item_id}'" in source or f'"{item_id}"' in source, (
                    f"{entry['entry_id']}: Tab 组件里找不到字面量键 {item_id}"
                )
            else:
                pytest.fail(f"{entry['entry_id']}: 未知 owner_declaration_kind={kind!r}")
            line_no = _line_no_of(table["owner_constant_source"])
            lines = source.splitlines()
            assert 0 < line_no <= len(lines), f"{entry['entry_id']}: owner_constant_source 行号越界"

    def test_primary_table_identity_cell_matches_the_authoritative_template(
        self, manifest_slice: dict
    ) -> None:
        for entry in manifest_slice["independent_entries"]:
            table = entry["html_counterpart"]["primary_table"]
            path = TEMPLATE_DIR / table["template_relative_path"]
            assert path.exists(), f"{entry['entry_id']}: 权威模板不存在 {path}"
            value = _cell_value(path, table["template_sheet"], table["identity_cell"])
            assert str(value).strip() == table["identity_text"], (
                f"{entry['entry_id']}: {table['template_sheet']}!{table['identity_cell']} 实读 "
                f"{value!r} != 声明 {table['identity_text']!r}"
            )

    def test_template_ref_resolves_through_the_runtime_index(self, manifest_slice: dict) -> None:
        indexed = _indexed_relpaths()
        for entry in manifest_slice["independent_entries"]:
            ref = entry["template_ref"]
            assert (TEMPLATE_DIR / ref).exists(), f"{entry['entry_id']}: template_ref 不在磁盘"
            assert ref in indexed, (
                f"{entry['entry_id']}: template_ref {ref!r} 不在 _index.json ⇒ 运行时不可达"
            )

    def test_measurement_variants_are_real_sheets_with_separate_owners(
        self, manifest_slice: dict
    ) -> None:
        """HD-5：双计量模式的两套 sheet 与两个 owner 模块都必须真实存在。"""
        found = 0
        for entry in manifest_slice["independent_entries"]:
            variant = entry["html_counterpart"].get("measurement_variant")
            if not variant:
                continue
            found += 1
            path = TEMPLATE_DIR / entry["template_ref"]
            sheets = _sheet_names(path)
            owners: set[str] = set()
            for item in variant["variants"]:
                assert item["template_sheet"] in sheets, (
                    f"{entry['entry_id']}: 模板里没有 sheet {item['template_sheet']!r}"
                )
                owner = _resolve_repo(item["owner_module"])
                assert owner.exists(), f"{entry['entry_id']}: owner_module 不存在 {item['owner_module']}"
                assert item["primary_table"] in owner.read_text(encoding="utf-8"), (
                    f"{entry['entry_id']}: {item['owner_module']} 里找不到键 {item['primary_table']}"
                )
                owners.add(item["owner_module"])
            assert len(owners) == len(variant["variants"]), (
                f"{entry['entry_id']}: 两个 variant 共用同一个 owner ⇒ 与声明不符"
            )
            assert _resolve_repo(variant["switch_source"]).exists()
        assert found == manifest_slice["honest_adjudication_summary"][
            "entries_with_dual_measurement_variants"
        ], "双计量模式 entry 数与摘要不符"

    def test_h7_second_write_path_is_real(self, manifest_slice: dict) -> None:
        """H7 是全 slice 唯一双写路径双客户端 —— 两条都必须可复核，且客户端确实不同。"""
        entry = next(
            e
            for e in manifest_slice["independent_entries"]
            if e["entry_id"] == "xlsx/gt-h7-biological-assets"
        )
        block = entry["html_counterpart"]["second_write_path"]
        clients: set[str] = set()
        for item in block["paths"]:
            path = _resolve_repo(item["path"])
            assert path.exists(), f"H7 second_write_path 文件不存在 {item['path']}"
            source = path.read_text(encoding="utf-8")
            assert _client_probe(item["client"], "put")(source), (
                f"H7 {item['path']} 里没有 {item['client']}.put(`...checklist-responses`)"
            )
            line_no = _line_no_of(item["put_source"])
            assert f"{item['client']}.put(" in source.splitlines()[line_no - 1]
            clients.add(item["client"])
        assert len(clients) == 2, f"H7 声明两个不同客户端，实际 {clients}"

    def test_h10_third_client_side_store_is_declared_and_unique(
        self, manifest_slice: dict
    ) -> None:
        """HD-6：只有 H10 的载体使用 localStorage 草稿，其余 8 条不得出现（防特例当通例）。"""
        for entry in manifest_slice["independent_entries"]:
            hc = entry["html_counterpart"]
            carrier = _resolve_repo(hc["write_carrier_path"])
            source = carrier.read_text(encoding="utf-8")
            extra = hc.get("extra_client_store")
            if extra:
                assert extra["kind"] == "localStorage_draft_fallback"
                assert "localStorage.setItem(draftKey(" in source, (
                    f"{entry['entry_id']}: 声明 localStorage 草稿，源码里找不到写入点"
                )
                assert "restoreDrafts" in source, (
                    f"{entry['entry_id']}: 声明草稿回灌，源码里找不到 restoreDrafts"
                )
            else:
                assert "draftKey(" not in source, (
                    f"{entry['entry_id']}: 载体里出现 draft 形态 localStorage 却未登记 extra_client_store"
                )


# ════════════════════════════════════════════════════════════════════════════
# Property 22：动态列 key 与 label 解耦
# ════════════════════════════════════════════════════════════════════════════
class TestProperty22DynamicColumnKeyDecoupling:
    """**Validates: Requirements 6.4**

    真分母双向验：正例 = H7 稳定 key 范式的 4 个结构要素；反例 = 全 H 循环
    label-as-key 的**穷举**命中必须恰好等于 slice 声明的 2 处。另断言硬编码列数/行数为 0。
    """

    def test_paradigm_source_declares_four_structural_requirements(
        self, manifest_slice: dict
    ) -> None:
        block = manifest_slice["dynamic_column_identity"]["paradigm_source"]
        ids = [r["id"] for r in block["structural_requirements"]]
        assert ids == ["SK-1", "SK-2", "SK-3", "SK-4"], f"结构要素登记不完整：{ids}"
        assert _resolve_repo(block["module"]).exists()

    def test_sk1_column_key_and_label_are_separate_fields(self) -> None:
        source = _strip_ts_comments(H7_LISTED_MODEL.read_text(encoding="utf-8"))
        match = re.search(r"export interface H7ListedCategory \{([\s\S]*?)\n\}", source)
        assert match, "找不到 H7ListedCategory 接口声明"
        body = match.group(1)
        assert re.search(r"^\s*key:\s*string", body, re.M), "接口里没有独立的 key 字段"
        assert re.search(r"^\s*label:\s*string", body, re.M), "接口里没有独立的 label 字段"
        assert re.search(r"createDefaultH7Categories", source), "找不到默认列构造函数"
        assert re.search(r"key:\s*`\$\{ind\.key\}_1`", source), (
            "默认列的 key 不是 `${ind.key}_1` ⇒ SK-1 的稳定前缀+序号形态不成立"
        )

    def test_sk2_next_key_takes_max_plus_one_and_never_reuses(self) -> None:
        for module in (H7_LISTED_MODEL, H7_SOE_MODEL):
            source = _strip_ts_comments(module.read_text(encoding="utf-8"))
            assert "`${prefix}${max + 1}`" in source, (
                f"{module.name}: 找不到 `${{prefix}}${{max + 1}}` ⇒ 序号可能复用已删值"
            )
            assert not re.search(r"\$\{\s*(?:\w+\.)?length\s*\+\s*1\s*\}", source), (
                f"{module.name}: 用 length + 1 生成序号 ⇒ 删中间项后会复用已删序号"
            )
            assert not re.search(r"`\$\{prefix\}\$\{\s*(?:i|idx|index)\s*\}`", source), (
                f"{module.name}: 用数组下标生成序号"
            )

    def test_sk3_total_column_reduces_over_the_dynamic_array(self) -> None:
        source = _strip_ts_comments(H7_LISTED_MODEL.read_text(encoding="utf-8"))
        assert re.search(r"categories\.reduce\(", source), (
            "合计列没有对 categories 数组 reduce ⇒ 列数可能被写死"
        )

    def test_sk4_row_model_is_a_declarative_array_not_fixed_blank_rows(
        self, h_files: list[pathlib.Path]
    ) -> None:
        source = _strip_ts_comments(H7_LISTED_MODEL.read_text(encoding="utf-8"))
        for name in ("H7_COST_MOVEMENT_ROWS", "H7_FAIR_MOVEMENT_ROWS"):
            assert re.search(rf"export const {name}\s*:\s*H7MovementRowDef\[\]", source), (
                f"找不到声明式行模型 {name}"
            )
        offenders = [
            f"{p.relative_to(ROOT).as_posix()}#L{no}"
            for p in h_files
            for no, line in enumerate(
                _strip_ts_comments(p.read_text(encoding="utf-8", errors="replace")).splitlines(), 1
            )
            if _BLANK_ROWS_FIXED.search(line)
        ]
        assert offenders == [], f"动态区骨架行数被写死：{offenders}"

    def test_label_as_key_hits_are_exactly_the_declared_deviations(
        self, manifest_slice: dict, h_files: list[pathlib.Path]
    ) -> None:
        """穷举等值：多一处 = 新增背离；少一处 = 已修好但登记未删。两个方向都打红。"""
        block = manifest_slice["dynamic_column_identity"]
        declared = block["hardcoded_scan_result"]["patterns"]
        hits = _label_key_hits(h_files)
        assert len(hits["column_key_is_label"]) == declared["column_key_is_label"], (
            f"column_key_is_label 实测 {hits['column_key_is_label']} 条，声明 "
            f"{declared['column_key_is_label']} 条"
        )
        assert len(hits["row_cell_key_is_label"]) == declared["row_cell_key_is_label"], (
            f"row_cell_key_is_label 实测 {hits['row_cell_key_is_label']} 条，声明 "
            f"{declared['row_cell_key_is_label']} 条"
        )
        deviations = block["deviations"]
        assert len(deviations) == 1, "声明的背离条数变了，判据需重写"
        module = "audit-platform/frontend/src/components/workpaper/composables/h8DisclosureSyncPayload.ts"
        for form in ("column_key_is_label", "row_cell_key_is_label"):
            assert all(h.startswith(module) for h in hits[form]), (
                f"{form} 的命中不全在 H8 同步载荷里：{hits[form]}"
            )
        assert deviations[0]["entry_id"] == "xlsx/gt-h8-right-of-use-assets"
        assert deviations[0]["registered_as"] == "BP-7"

    def test_no_hardcoded_horizontal_company_columns(
        self, manifest_slice: dict, h_files: list[pathlib.Path]
    ) -> None:
        declared = manifest_slice["dynamic_column_identity"]["hardcoded_scan_result"]["patterns"]
        offenders: list[str] = []
        for path in h_files:
            body = _strip_ts_comments(path.read_text(encoding="utf-8", errors="replace"))
            for no, line in enumerate(body.splitlines(), 1):
                if _HORIZONTAL_COLUMN_LITERAL.search(line):
                    offenders.append(f"{path.relative_to(ROOT).as_posix()}#L{no}")
        assert len(offenders) == declared["horizontal_company_column_literals"], (
            f"横向展开列名字面量实测 {offenders}，声明 "
            f"{declared['horizontal_company_column_literals']} 条"
        )

    def test_h8_deviation_is_reachable_from_user_input(self, manifest_slice: dict) -> None:
        """BP-7 的后果链必须可复核：类别 label 由用户输入 ⇒ 重名可达 ⇒ 撞键可达。"""
        source = (COMPOSABLES / "useH8Disclosure.ts").read_text(encoding="utf-8")
        assert re.search(r"function addCategory\(\s*label:\s*string\s*\)", source), (
            "useH8Disclosure.addCategory 的签名变了 —— BP-7 的可达性论证要重写"
        )
        assert "`cat_${Date.now().toString(36)}`" in source, (
            "底稿内 key 不再是稳定串 ⇒ BP-7 的 scope_note（只影响同步边界）不再成立"
        )
        payload = H8_SYNC_PAYLOAD.read_text(encoding="utf-8")
        assert "h8ListedCellValue(state.movement, def, c.key)" in payload, (
            "取值端不再按 c.key ⇒ BP-7 的「取值稳定但键不稳定」论证要重写"
        )


# ════════════════════════════════════════════════════════════════════════════
# Property 23：动态行身份不使用下标
# ════════════════════════════════════════════════════════════════════════════
class TestProperty23DynamicRowIdentity:
    """**Validates: Requirements 6.5**

    9 个 primary table 的行身份构造点逐条 source-backed；位置化命中穷举分三族且与
    slice 声明逐条等值；反向自检 3 个展示序号站点**不得**被点名。
    """

    def test_every_entry_declares_its_identity_key_and_generator(
        self, manifest_slice: dict
    ) -> None:
        keys: set[str] = set()
        for entry in manifest_slice["independent_entries"]:
            hc = entry["html_counterpart"]
            key = hc["row_identity_key"]
            keys.add(key)
            assert key in _IDENTITY_KEYS, f"{entry['entry_id']}: 未知身份字段 {key!r}"
            path = _resolve_repo(hc["row_identity_generator_source"])
            assert path.exists(), f"{entry['entry_id']}: 生成器源文件不存在"
            line_no = _line_no_of(hc["row_identity_generator_source"])
            lines = path.read_text(encoding="utf-8").splitlines()
            assert 0 < line_no <= len(lines), f"{entry['entry_id']}: 生成器行号越界"
            expr = _value_expr_after_key(lines[line_no - 1], key)
            if expr is None:
                # H8 的 `_id()` 是独立函数体的 return，不是对象属性 —— 退到整行核对
                assert "Date.now()" in lines[line_no - 1] or "Math.random()" in lines[line_no - 1], (
                    f"{entry['entry_id']}: 生成器行既不是 `{key}:` 赋值也不含随机源："
                    f"{lines[line_no - 1].strip()[:100]!r}"
                )
            else:
                assert "Date.now()" in expr or "Math.random()" in expr or "(" in expr, (
                    f"{entry['entry_id']}: `{key}` 的值表达式看不出生成来源：{expr.strip()[:100]!r}"
                )
        assert keys == {"rowId", "id"}, (
            f"H 循环实测身份字段两族（rowId 八条 / id 一条），当前 {keys} —— "
            "判据写死任一族都会把另一族判错"
        )

    def test_positional_identity_inventory_is_exhaustive_and_partitioned(
        self, manifest_slice: dict, h_files: list[pathlib.Path]
    ) -> None:
        inventory = manifest_slice["dynamic_row_identity"]["positional_identity_inventory"]
        hits = _positional_identity_hits(h_files)
        assert len(hits) == inventory["total_hits"], (
            f"位置化命中实测 {len(hits)} 条，slice 声明 {inventory['total_hits']} 条：\n"
            + "\n".join(f"  {r}#L{n} {k}:{e[:70]}" for r, n, k, e in hits)
        )
        family_a = {h["source_ref"] for h in inventory["family_a_true_defect_primary_table_seed"]["hits"]}
        family_b = set(inventory["family_b_index_only_as_prefix_into_a_random_generator"]["hits"])
        family_c = set(
            inventory["family_c_index_as_fallback_on_non_primary_derived_tables"]["hits"]
        )
        assert len(family_a) == inventory["family_a_true_defect_primary_table_seed"]["count"]
        assert len(family_b) == inventory[
            "family_b_index_only_as_prefix_into_a_random_generator"
        ]["count"]
        assert len(family_c) == inventory[
            "family_c_index_as_fallback_on_non_primary_derived_tables"
        ]["count"]
        assert not (family_a & family_b) and not (family_a & family_c) and not (family_b & family_c), (
            "三族有重叠 ⇒ 分类不是划分"
        )
        declared = family_a | family_b | family_c
        actual = {f"{rel}#L{no}" for rel, no, _k, _e in hits}
        assert declared == actual, (
            "位置化清单与实测集合不等：\n"
            f"  仅在声明里：{sorted(declared - actual)}\n"
            f"  仅在实测里：{sorted(actual - declared)}"
        )

    def test_family_a_hits_write_to_the_declared_key(self, manifest_slice: dict) -> None:
        """构造点与写入点可能在不同文件（H4 就是），故按 writes_to_key_site 分别核。"""
        inventory = manifest_slice["dynamic_row_identity"]["positional_identity_inventory"]
        entry_ids = {e["entry_id"] for e in manifest_slice["independent_entries"]}
        for hit in inventory["family_a_true_defect_primary_table_seed"]["hits"]:
            source = _resolve_repo(hit["source_ref"]).read_text(encoding="utf-8")
            line_no = _line_no_of(hit["source_ref"])
            expr = _value_expr_after_key(source.splitlines()[line_no - 1], "rowId")
            assert expr is not None and _POSITIONAL_IDENTITY_TOKEN.search(expr), (
                f"{hit['source_ref']}: 声明为位置化，实测其 rowId 表达式不是"
            )
            assert hit["entry_id"] in entry_ids, f"{hit['source_ref']}: entry_id 不属本 slice"
            key = hit["writes_to_key"]
            site = _resolve_repo(hit["writes_to_key_site"])
            site_source = site.read_text(encoding="utf-8")
            assert f"'{key}'" in site_source or f'"{key}"' in site_source, (
                f"{hit['source_ref']}: 写入点 {hit['writes_to_key_site']} 里找不到键 {key!r}"
            )
            if "#L" in hit["writes_to_key_site"]:
                site_line = site_source.splitlines()[_line_no_of(hit["writes_to_key_site"]) - 1]
                assert key in site_line, (
                    f"{hit['source_ref']}: writes_to_key_site 指向的行不含键 {key!r}"
                    f"（实际 {site_line.strip()[:90]!r}）"
                )
        assert inventory["family_a_true_defect_primary_table_seed"]["registered_as"] == "BP-6"

    def test_display_sequence_sites_are_not_flagged(self, manifest_slice: dict) -> None:
        """反向自检：`seq: raw.seq ?? idx + 1` 三处**不得**被身份判据点名。"""
        guard = manifest_slice["dynamic_row_identity"]["positional_identity_inventory"][
            "false_positive_guard"
        ]
        for ref in guard["must_not_be_flagged"]:
            path = _resolve_repo(ref)
            line_no = _line_no_of(ref)
            line = path.read_text(encoding="utf-8").splitlines()[line_no - 1]
            assert "seq" in line, f"{ref}: 该行已不是展示序号站点，自检前提失效"
            for key in _IDENTITY_KEYS:
                expr = _value_expr_after_key(line, key)
                assert expr is None or not _POSITIONAL_IDENTITY_TOKEN.search(expr), (
                    f"{ref}: 展示序号行被身份判据误判成位置化（key={key}）"
                )
            seq_expr = _value_expr_after_key(line, "seq")
            assert seq_expr is not None and _POSITIONAL_IDENTITY_TOKEN.search(seq_expr), (
                f"{ref}: 该行的 seq 表达式已不含位置 token，自检前提失效"
            )

    def test_dynamic_row_identity_tables_cover_every_entry(self, manifest_slice: dict) -> None:
        tables = manifest_slice["dynamic_row_identity"]["tables"]
        assert len(tables) == len(manifest_slice["independent_entries"])
        by_entry = {t["entry_id"]: t for t in tables}
        forbidden = manifest_slice["dynamic_row_identity"]["forbidden_identity_kinds"]
        for entry in manifest_slice["independent_entries"]:
            table = by_entry[entry["entry_id"]]
            assert table["table_key"] == entry["html_counterpart"]["primary_table"]["item_id"]
            identity = table["row_identity"]
            assert identity["identity_field"] == entry["html_counterpart"]["row_identity_key"]
            assert identity["kind"] not in forbidden, (
                f"{entry['entry_id']}: row_identity.kind 落在 forbidden_identity_kinds 里"
            )
            assert _resolve_repo(identity["source_ref"]).exists()

    def test_positional_seed_entries_match_the_summary_counter(self, manifest_slice: dict) -> None:
        actual = sum(
            1
            for e in manifest_slice["independent_entries"]
            if "prefill_array_index_seed" in e["html_counterpart"]["primary_table"].get("item_id", "")
            or "array_index_seed"
            in str(
                next(
                    t["row_identity"]["kind"]
                    for t in manifest_slice["dynamic_row_identity"]["tables"]
                    if t["entry_id"] == e["entry_id"]
                )
            )
        )
        assert actual == manifest_slice["honest_adjudication_summary"][
            "entries_with_positional_seed_row_identity"
        ], "带位置化种子的 entry 数与摘要不符"


# ════════════════════════════════════════════════════════════════════════════
# BP-5：H8 四表种子写进零消费键
# ════════════════════════════════════════════════════════════════════════════
class TestOrphanSeedKey:
    """**Validates: Requirements 12.11, 14.1**

    「某个键零消费」这件事必须现扫复算 —— 写死「只有 1 处」会在有人接上消费方之后仍绿。
    """

    def test_h8_seed_key_has_exactly_one_occurrence_in_the_repo(
        self, manifest_slice: dict
    ) -> None:
        entry = next(
            e
            for e in manifest_slice["independent_entries"]
            if e["entry_id"] == "xlsx/gt-h8-right-of-use-assets"
        )
        block = entry["html_counterpart"]["orphan_seed_key"]
        key = block["key"]
        occurrences: list[str] = []
        for base, exts in ((FRONTEND, (".ts", ".vue")), (BACKEND / "app", (".py",))):
            for path in base.rglob("*"):
                if not (path.is_file() and path.suffix in exts):
                    continue
                text = path.read_text(encoding="utf-8", errors="replace")
                if key in text:
                    for no, line in enumerate(text.splitlines(), 1):
                        if key in line:
                            occurrences.append(f"{path.relative_to(ROOT).as_posix()}#L{no}")
        assert len(occurrences) == 1, (
            f"{key!r} 在全仓命中 {len(occurrences)} 处 —— 若已有消费方，BP-5 该解除并删除登记："
            f"{occurrences}"
        )
        assert occurrences[0] == block["write_site"], (
            f"唯一命中 {occurrences[0]} 与声明的写入点 {block['write_site']} 不一致"
        )
        assert block["read_sites"] == []

    def test_h8_real_primary_key_is_a_different_literal(self, manifest_slice: dict) -> None:
        entry = next(
            e
            for e in manifest_slice["independent_entries"]
            if e["entry_id"] == "xlsx/gt-h8-right-of-use-assets"
        )
        real = entry["html_counterpart"]["primary_table"]["item_id"]
        orphan = entry["html_counterpart"]["orphan_seed_key"]["key"]
        assert real != orphan, "BP-5 的前提（种子键 != 主表键）不成立"
        owner = _resolve_repo(entry["html_counterpart"]["primary_table"]["owner_module"])
        assert f"'{real}'" in owner.read_text(encoding="utf-8")

    def test_backend_preset_anchor_points_at_the_real_key_not_the_orphan(self) -> None:
        """BP-5 的可观察后果之一：后端登记的取数锚点与前端实际写入的键不一致。"""
        presets = (BACKEND / "app" / "services" / "d_cycle_extraction" / "presets.py").read_text(
            encoding="utf-8"
        )
        assert '"anchor": "H8-2-rows"' in presets, (
            "后端 preset 的 H8 锚点变了 —— BP-5 的 observable_consequences 第 2 条要重写"
        )
        assert "H8-2-detail-prefill" not in presets


# ════════════════════════════════════════════════════════════════════════════
# BP-8：孤儿 / 错位 legacy composable
# ════════════════════════════════════════════════════════════════════════════
class TestOrphanLegacyComposables:
    """**Validates: Requirements 1.7, 12.11**

    两侧都断言：声称零消费的真零、声称有消费方的真有。**只断言一侧会 fail-open** ——
    只查「零消费的真零」时，某天有人接上消费方也不会红；只查「有消费方的真有」时，
    孤儿登记漏掉一个也不会红。
    """

    def test_declared_dual_mode_consumers_match_the_source(self, manifest_slice: dict) -> None:
        orphan_count = 0
        for entry in manifest_slice["independent_entries"]:
            legacy = entry["legacy_dual_mode"]
            module = _resolve_repo(legacy["module"])
            assert module.exists(), f"{entry['entry_id']}: legacy composable 不存在"
            actual = _import_specifier_consumers(module.stem)
            production = [c for c in actual if "__tests__" not in c and not c.endswith(".spec.ts")]
            assert sorted(legacy["consumers"]) == sorted(production), (
                f"{entry['entry_id']}: {module.name} 的生产消费方声明 {legacy['consumers']} "
                f"!= 实测 {production}"
            )
            host_rel = entry["host_path"]
            assert legacy["host_consumes_it"] == (host_rel in production), (
                f"{entry['entry_id']}: host_consumes_it 与实测不符"
            )
            if not production:
                orphan_count += 1
            assert legacy["wraps_shared_base"] is False, (
                f"{entry['entry_id']}: 声明不 wraps 共享基类，与 slice 的假设不符"
            )
            assert "useWorkpaperEntryDualMode" not in module.read_text(encoding="utf-8"), (
                f"{entry['entry_id']}: {module.name} 实际 import 了共享基类"
            )
            prefix = legacy["localStorage_prefix"]
            assert f"'{prefix}'" in module.read_text(encoding="utf-8"), (
                f"{entry['entry_id']}: 找不到声明的 localStorage 前缀 {prefix!r}"
            )
        assert orphan_count == 3, (
            f"零消费的 dual-mode composable 实测 {orphan_count} 个（H5/H7/H9），"
            "变了说明 BP-8 的分母漂移，需更新登记"
        )

    def test_declared_orphan_formdata_composables_really_have_no_production_consumer(
        self, manifest_slice: dict
    ) -> None:
        declared = 0
        for entry in manifest_slice["independent_entries"]:
            block = entry.get("orphan_formdata_composable")
            if not block:
                continue
            declared += 1
            module = _resolve_repo(block["module"])
            assert module.exists(), f"{entry['entry_id']}: {block['module']} 不存在"
            actual = _import_specifier_consumers(module.stem)
            production = [c for c in actual if "__tests__" not in c and not c.endswith(".spec.ts")]
            tests = [c for c in actual if c not in production]
            assert production == block["production_consumers"] == [], (
                f"{entry['entry_id']}: {module.name} 声称生产零消费，实测 {production}"
                " ⇒ BP-8 该解除并删除登记"
            )
            assert sorted(tests) == sorted(block["test_only_consumers"]), (
                f"{entry['entry_id']}: {module.name} 的测试消费方声明与实测不符：{tests}"
            )
            declared_count = manifest_slice["honest_adjudication_summary"][
                "entries_with_orphan_legacy_composable"
            ]
            assert declared_count == 5, "带孤儿 legacy 载体的 entry 数登记变了"
        assert declared == 4, (
            f"登记了 {declared} 条 orphan_formdata_composable（应为 H6/H7/H8/H9 四条）"
        )

    def test_non_orphan_formdata_composables_are_not_registered_as_orphans(
        self, manifest_slice: dict
    ) -> None:
        """反向：H3/H4/H5/H10 的 FormData 是真载体，不得被登记成孤儿。"""
        real_carriers = {
            "useH3FormData",
            "useH4FormData",
            "useH5FormData",
            "useH10FormData",
        }
        registered = {
            _resolve_repo(e["orphan_formdata_composable"]["module"]).stem
            for e in manifest_slice["independent_entries"]
            if e.get("orphan_formdata_composable")
        }
        assert not (registered & real_carriers), (
            f"真载体被误登记成孤儿：{sorted(registered & real_carriers)}"
        )
        for stem in real_carriers:
            production = [
                c
                for c in _import_specifier_consumers(stem)
                if "__tests__" not in c and not c.endswith(".spec.ts")
            ]
            assert production, f"{stem} 实测零生产消费 ⇒ 它也该登记成孤儿"

    def test_shared_base_is_preserved_with_its_real_consumer_count(
        self, deletion_plan: dict
    ) -> None:
        block = deletion_plan["shared_base_preserved"]
        base = _resolve_repo(block["file"])
        assert base.exists()
        actual = _import_specifier_consumers(base.stem)
        assert len(actual) == block["remaining_consumers_after_h_cycle"], (
            f"共享基类消费方实测 {len(actual)} 个，声明 "
            f"{block['remaining_consumers_after_h_cycle']} 个"
        )
        h_consumers = [c for c in actual if re.search(r"[Hh](?:2|3|4|5|6|7|8|9|10)", c)]
        assert not any("DualMode" in c for c in h_consumers), (
            f"H 循环的 dual-mode composable 实际 wraps 了共享基类：{h_consumers}"
        )

    def test_host_inlined_second_implementation_is_real(self, deletion_plan: dict) -> None:
        """BP-8 的另一半：4 个宿主内联的真载体必须逐行可复核。"""
        block = deletion_plan["host_inlined_second_implementation"]
        assert len(block["hosts"]) == 4
        for host in block["hosts"]:
            path = _resolve_repo(host["host_path"])
            assert path.exists()
            lines = path.read_text(encoding="utf-8").splitlines()
            state_line = lines[_line_no_of(host["inline_mode_state"]) - 1]
            assert "currentMode" in state_line and "ref<" in state_line, (
                f"{host['entry_id']}: inline_mode_state 指向的行不是 currentMode ref："
                f"{state_line.strip()[:90]!r}"
            )
            if host["inline_switch_fn"]:
                fn_line = lines[_line_no_of(host["inline_switch_fn"]) - 1]
                assert "switchMode" in fn_line or "currentMode.value" in fn_line, (
                    f"{host['entry_id']}: inline_switch_fn 指向的行不是切换逻辑"
                )


# ════════════════════════════════════════════════════════════════════════════
# Property 20 / 21：分母为空，不宣称通过
# ════════════════════════════════════════════════════════════════════════════
class TestProperty20And21NotClaimed:
    """**Validates: Requirements 12.1**

    本 slice 的 per-entry contract 数为 0（唯一一份 H 循环契约属已排除的 H1 pilot）⇒
    Property 20 / 21 的分母为空，**不宣称通过**。只断言两件可复核的事：前提成立 +
    承载者存在。
    """

    def test_no_slice_entry_has_a_contract_and_h1_is_the_only_h_contract(
        self, manifest_slice: dict
    ) -> None:
        slice_ids = {e["entry_id"] for e in manifest_slice["independent_entries"]}
        h_contracts: list[str] = []
        for path in sorted(CONTRACT_DIR.glob("*.json")):
            doc = _load(path)
            owner = str((doc.get("review") or {}).get("entry_id") or "")
            assert owner not in slice_ids, (
                f"{path.name} 的 review.entry_id={owner!r} 属本 slice ⇒ Property 21 的空分母"
                "前提不再成立，必须在此补齐字段级判据"
            )
            if owner.startswith("xlsx/gt-h"):
                h_contracts.append(f"{path.name}:{owner}")
        assert h_contracts == [f"h1.disposal_check.json:{H1_ENTRY_ID}"], (
            f"H 循环契约集合实测 {h_contracts} —— 与「唯一一份属 H1 pilot」的前提不符"
        )

    def test_property_21_carriers_exist(self) -> None:
        for name in (
            "test_task13_contract_registry.py",
            "test_task42_h1_grouped_dynamic_pilot.py",
        ):
            assert (_THIS.parent / name).exists(), f"缺 Property 21 的字段级判据承载者 {name}"

    def test_no_slice_entry_has_a_registered_adapter(
        self, manifest_slice: dict, full_manifest: dict
    ) -> None:
        by_id = {e["entry_id"]: e for e in full_manifest["entries"]}
        for entry in manifest_slice["independent_entries"]:
            assert entry["adapter_id"] is None
            assert by_id[entry["entry_id"]]["adapter_id"] is None, (
                f"{entry['entry_id']}: source manifest 里已有 adapter_id ⇒ slice 该更新"
            )

    def test_registry_delivered_contracts_contain_no_slice_entry(
        self, manifest_slice: dict
    ) -> None:
        registry = REGISTRY.read_text(encoding="utf-8")
        for entry in manifest_slice["independent_entries"]:
            assert entry["entry_id"] not in registry, (
                f"{entry['entry_id']}: 出现在 adapters/registry.py 里 ⇒ 与 adapter_id=None 矛盾"
            )
        assert H1_ENTRY_ID in registry, (
            "registry 里找不到 H1 pilot ⇒ 「同循环 pilot 已交付、本 slice 未交付」这个对照消失"
        )

    def test_property_denominator_block_declares_what_is_not_claimed(
        self, manifest_slice: dict
    ) -> None:
        block = manifest_slice["property_denominators"]
        assert block["property_20_and_21_not_claimed"]["not_claimed_passing"] is True
        assert block["property_22"]["not_claimed_passing"] is False
        assert block["property_23"]["not_claimed_passing"] is False
        assert block["property_69"]["not_claimed_passing_part"]
        assert block["property_70"]["not_claimed_passing"] is False
        for key in ("property_22", "property_23", "property_69", "property_70"):
            assert "h_cycle_denominator" in block[key]
            assert "how_handled" in block[key]


# ════════════════════════════════════════════════════════════════════════════
# Property 28：immutable definition 漂移 fail closed（真分母）
# ════════════════════════════════════════════════════════════════════════════
class TestProperty28DefinitionDriftFailClosed:
    """**Validates: Requirements 6.10, 12.4**

    真分母真验：11 个模板 size+sha256 现算、9 个 normalized_structure_hash 现算、
    9 组 sheet!cell 真读（在 TestHtmlCounterpart 里）、21 条模板解析实跑。
    """

    def test_authoritative_templates_digests_recompute(self, manifest_slice: dict) -> None:
        block = manifest_slice["authoritative_templates"]
        root = ROOT / block["root"]
        assert root.is_dir(), f"权威目录不存在 {root}"
        on_disk = sorted(p.name for p in root.iterdir() if p.is_file() and not p.name.startswith("~$"))
        declared = sorted(f["name"] for f in block["files"])
        assert declared == on_disk, (
            f"登记集合 != 磁盘实况\n  仅登记：{sorted(set(declared) - set(on_disk))}\n"
            f"  仅磁盘：{sorted(set(on_disk) - set(declared))}"
        )
        for item in block["files"]:
            path = root / item["name"]
            assert path.stat().st_size == item["size"], f"{item['name']}: size 漂移"
            assert _sha256_of(path) == item["sha256"], f"{item['name']}: sha256 漂移"
            assert re.fullmatch(r"[0-9a-f]{64}", item["sha256"]), f"{item['name']}: 不是真 digest"

    def test_normalized_structure_hash_recomputes(self, manifest_slice: dict) -> None:
        from app.services.workpaper_sync.excel_instrumentation import normalized_structure_hash

        block = manifest_slice["authoritative_templates"]
        root = ROOT / block["root"]
        count = 0
        for item in block["files"]:
            if "normalized_structure_hash" not in item:
                continue
            count += 1
            actual = normalized_structure_hash((root / item["name"]).read_bytes())
            assert actual == item["normalized_structure_hash"], (
                f"{item['name']}: normalized_structure_hash 漂移（实算 {actual}）"
            )
        assert count == len(manifest_slice["independent_entries"]), (
            f"只有 {count} 个模板带 normalized_structure_hash，应为 9 个（每 entry 一本）"
        )

    def test_template_owner_mapping_is_a_bijection_over_the_entries(
        self, manifest_slice: dict
    ) -> None:
        """H 循环一册对一 entry ⇒ 可比 F/G 更严：既单射也满射。"""
        owned = [
            f["belongs_to_entry"]
            for f in manifest_slice["authoritative_templates"]["files"]
            if f.get("belongs_to_entry")
        ]
        entry_ids = [e["entry_id"] for e in manifest_slice["independent_entries"]]
        assert len(owned) == len(set(owned)), f"belongs_to_entry 有重复（非单射）：{owned}"
        assert set(owned) == set(entry_ids), (
            f"模板归属不满射：无主的 entry {sorted(set(entry_ids) - set(owned))}"
        )
        by_entry = {f["belongs_to_entry"]: f["name"] for f in
                    manifest_slice["authoritative_templates"]["files"]
                    if f.get("belongs_to_entry")}
        for entry in manifest_slice["independent_entries"]:
            assert entry["template_ref"] == f"H/{by_entry[entry['entry_id']]}", (
                f"{entry['entry_id']}: template_ref 与 belongs_to_entry 不一致"
            )
        for item in manifest_slice["authoritative_templates"]["files"]:
            if item.get("belongs_to_entry") is None:
                assert item.get("excluded_reason"), f"{item['name']}: 无主模板缺 excluded_reason"

    def test_in_runtime_index_flag_recomputes_and_h_has_no_unreachable_workbook(
        self, manifest_slice: dict
    ) -> None:
        indexed = _indexed_relpaths()
        block = manifest_slice["authoritative_templates"]
        for item in block["files"]:
            rel = f"H/{item['name']}"
            assert item["in_runtime_index"] == (rel in indexed), (
                f"{item['name']}: in_runtime_index 声明与 _index.json 不符"
            )
        assert all(f["in_runtime_index"] for f in block["files"]), (
            "H 目录出现不在索引里的文件 ⇒ runtime_index_note 里「11 个文件全部在索引里」"
            "这句已失效，需要按 F 的 BP-8 形态新增阻断项"
        )

    def test_template_resolution_audit_recomputes(self, manifest_slice: dict) -> None:
        from app.services.wp_template_finder import find_template_file, find_template_file_any

        audit = manifest_slice["template_resolution_audit"]
        assert audit["result"] == "clean"
        for item in audit["measured"]:
            code = item["wp_code"]
            base = code.split("-")[0]
            for fn in (find_template_file, find_template_file_any):
                resolved = fn(code)
                assert resolved is not None, f"{code}: {fn.__name__} 返回 None"
                name = pathlib.Path(str(resolved)).name
                assert name == item["resolved"], (
                    f"{code}: {fn.__name__} 实解析 {name!r} != 声明 {item['resolved']!r}"
                )
                assert name.startswith(base + " "), (
                    f"{code}: 解析结果 {name!r} 不属 {base} 那本工作簿 ⇒ 出现 F2 那种回落"
                )
        assert find_template_file("H2A") is None, (
            "H2A 这类程序表码不该解析出独立模板 —— program_table_code_note 的前提失效"
        )


# ════════════════════════════════════════════════════════════════════════════
# Property 69：evidence 与计数（负向分母真验）
# ════════════════════════════════════════════════════════════════════════════
class TestProperty69EvidenceAndCounters:
    """**Validates: Requirements 12.10, 12.11, 12.12**

    slice_scope 按 selection_rule 从 source manifest 现算；summary 的 18 个计数逐项现算；
    evidence 的 UNVERIFIABLE 蕴含关系逐条断言。逐 scenario 正向闭合**不宣称通过**（BP-4）。
    """

    def test_slice_scope_is_recomputable_from_the_manifest(
        self, manifest_slice: dict, full_manifest: dict
    ) -> None:
        h_prefixed = [
            e
            for e in full_manifest["entries"]
            if any(str(p).startswith("H") for p in (e.get("wp_match") or {}).get("wp_code_patterns") or [])
        ]
        scope = manifest_slice["slice_scope"]
        assert len(h_prefixed) == scope["h_prefixed_entry_total"], (
            f"H 前缀 entry 实算 {len(h_prefixed)} 条，声明 {scope['h_prefixed_entry_total']} 条"
        )
        assert all(e["document_type"] == scope["document_type"] for e in h_prefixed), (
            "有 H 前缀 entry 的 document_type 不是 xlsx ⇒ selection_rule 第 ① 条要重写"
        )
        independent = [e for e in h_prefixed if e["independent_entry"]]
        assert len(independent) == scope["h_prefixed_independent_total"]
        children = [e for e in h_prefixed if not e["independent_entry"]]
        assert len(children) == scope["parent_duplicate_count"]
        expected = {e["entry_id"] for e in independent} - {H1_ENTRY_ID}
        actual = {e["entry_id"] for e in manifest_slice["independent_entries"]}
        assert expected == actual, (
            "slice 的 entry 集合与 selection_rule 现算不等：\n"
            f"  漏迁：{sorted(expected - actual)}\n  凑数：{sorted(actual - expected)}"
        )
        assert scope["independent_entry_count"] == len(actual) == 9
        assert scope["excluded_pilot_entry_count"] == 1

    def test_no_h_docx_entry_exists_in_the_source_manifest(self, full_manifest: dict) -> None:
        offenders = [
            e["entry_id"]
            for e in full_manifest["entries"]
            if e["document_type"] == "docx"
            and any(
                str(p).startswith("H")
                for p in (e.get("wp_match") or {}).get("wp_code_patterns") or []
            )
        ]
        assert offenders == [], (
            f"H 循环出现 docx entry {offenders} ⇒ selection_rule 的「H 循环无 Word 通道」失效"
        )

    def test_parent_duplicate_children_are_registered_and_disjoint(
        self, manifest_slice: dict, full_manifest: dict
    ) -> None:
        summary = manifest_slice["parent_duplicate_summary"]
        declared = {c["entry_id"] for c in summary["children"]}
        actual = {
            e["entry_id"]
            for e in full_manifest["entries"]
            if not e["independent_entry"]
            and any(
                str(p).startswith("H")
                for p in (e.get("wp_match") or {}).get("wp_code_patterns") or []
            )
        }
        assert declared == actual, (
            f"parent_duplicate 子入口登记与 manifest 现算不等：仅登记 {sorted(declared - actual)}，"
            f"仅 manifest {sorted(actual - declared)}"
        )
        entry_ids = {e["entry_id"] for e in manifest_slice["independent_entries"]}
        assert not (declared & entry_ids), "子入口混进了 independent_entries（违反 AC 1.6）"
        by_id = {e["entry_id"]: e for e in full_manifest["entries"]}
        for child in summary["children"]:
            assert by_id[child["entry_id"]]["parent_entry_id"] == child["parent_entry_id"]
            assert child["parent_entry_id"] in entry_ids
            assert _resolve_repo(child["host_path"]).exists()
            consumed = child.get("consumes_parent_legacy_composable")
            if consumed:
                assert _resolve_repo(consumed).exists()
                assert pathlib.Path(consumed).stem in _resolve_repo(child["host_path"]).read_text(
                    encoding="utf-8"
                ), f"{child['entry_id']}: 声明消费 {consumed}，源码里找不到该 import"

    def test_excluded_items_have_reasons(self, manifest_slice: dict) -> None:
        excluded = manifest_slice["slice_scope"]["excluded_from_slice"]
        assert len(excluded) >= 4
        for item in excluded:
            assert item.get("what") and item.get("reason")
            assert len(item["reason"]) > 40, f"排除理由过短：{item['what']}"

    def test_each_entry_has_unique_id(self, manifest_slice: dict) -> None:
        ids = [e["entry_id"] for e in manifest_slice["independent_entries"]]
        assert len(ids) == len(set(ids))

    def test_evidence_state_and_reasons(self, manifest_slice: dict) -> None:
        for entry in manifest_slice["independent_entries"]:
            ev = entry["evidence"]
            assert ev["verification_state"] == "UNVERIFIABLE"
            reasons = ev.get("unverifiable_reasons")
            assert isinstance(reasons, list) and reasons, (
                f"{entry['entry_id']}: UNVERIFIABLE 必须带非空 unverifiable_reasons（SR-7）"
            )
            assert len(reasons) == len(set(reasons))

    def test_contract_test_points_at_this_file(self, manifest_slice: dict) -> None:
        expected = _THIS.relative_to(ROOT).as_posix()
        for entry in manifest_slice["independent_entries"]:
            assert entry["evidence"]["contract_test"] == expected, (
                f"{entry['entry_id']}: contract_test 未指向本文件"
            )
            assert entry["dom_guard_ref"].startswith(expected + "::")

    def test_no_entry_claims_verified_without_a_test_run(self, manifest_slice: dict) -> None:
        for entry in manifest_slice["independent_entries"]:
            ev = entry["evidence"]
            if ev["verification_state"] == "VERIFIED":
                assert ev["sync_test_run_id"] and ev["required_scenario_set_digest"]
            else:
                assert ev["sync_test_run_id"] is None
                assert ev["required_scenario_set_digest"] is None

    def test_summary_counters_are_recomputed_from_entries(self, manifest_slice: dict) -> None:
        entries = manifest_slice["independent_entries"]
        summary = manifest_slice["honest_adjudication_summary"]
        expected = {
            "total_independent": len(entries),
            "adjudicated_as_bidirectional": sum(1 for e in entries if e["capability"] == "bidirectional"),
            "adjudicated_as_single_onlyoffice": sum(
                1 for e in entries if e["capability"] == "single_onlyoffice"
            ),
            "adjudicated_as_single_html": sum(1 for e in entries if e["capability"] == "single_html"),
            "adjudicated_as_unreachable": sum(1 for e in entries if e["capability"] == "unreachable"),
            "capability_verdict_pending": sum(1 for e in entries if e["capability"] is None),
            "html_counterpart_exists": sum(
                1 for e in entries if e["html_counterpart_verdict"] == "exists"
            ),
            "html_counterpart_none": sum(
                1 for e in entries if e["html_counterpart_verdict"] == "none"
            ),
            "finalized_published_representations": sum(
                1 for e in entries if e["published_representation"] is not None
            ),
            "entries_left_unverifiable": sum(
                1 for e in entries if e["evidence"]["verification_state"] == "UNVERIFIABLE"
            ),
            "parent_duplicate_children": len(manifest_slice["parent_duplicate_summary"]["children"]),
            "entries_with_host_inline_write_carrier": sum(
                1 for e in entries if e["html_counterpart"]["write_carrier"] == "host_inline"
            ),
            "entries_with_formdata_composable_write_carrier": sum(
                1 for e in entries if e["html_counterpart"]["write_carrier"] == "formdata_composable"
            ),
            "entries_with_per_tab_write_carrier": sum(
                1
                for e in entries
                if e["html_counterpart"]["write_carrier"].startswith("per_tab")
            ),
            "entries_without_checklist_get_in_their_carrier": sum(
                1
                for e in entries
                if e["html_counterpart"]["read_carrier"]
                in ("render_config_snapshot_passthrough", "host_inline_render_config_refetch")
            ),
            "entries_with_dual_measurement_variants": sum(
                1 for e in entries if e["html_counterpart"].get("measurement_variant")
            ),
            "entries_with_orphan_legacy_composable": sum(
                1
                for e in entries
                if e.get("orphan_formdata_composable")
                or not e["legacy_dual_mode"]["host_consumes_it"]
            ),
            "authoritative_template_files": len(
                manifest_slice["authoritative_templates"]["files"]
            ),
        }
        for key, value in expected.items():
            assert summary[key] == value, f"{key}: 声明 {summary[key]} != 现算 {value}"
        assert summary["pilot_contract_published"] == 0
        counters = summary["slice_counters"]
        assert counters["unadjudicated"] == expected["capability_verdict_pending"]
        assert counters["fake_bidirectional_claimed_verified"] == sum(
            1
            for e in entries
            if e["migration_state"] == "legacy_fake_bidirectional"
            and e["evidence"]["verification_state"] == "VERIFIED"
        )
        assert counters["bidirectional_unverified"] == sum(
            1
            for e in entries
            if e["capability"] == "bidirectional"
            and e["evidence"]["verification_state"] != "VERIFIED"
        )
        assert counters["stale_evidence"] == sum(
            1 for e in entries if e["evidence"]["required_scenario_set_digest"] is not None
        )

    def test_form_difference_block_is_complete(self, manifest_slice: dict) -> None:
        block = manifest_slice["h_cycle_form_differences"]
        ids = [d["id"] for d in block["differences"]]
        assert ids == ["HD-1", "HD-2", "HD-3", "HD-4", "HD-5", "HD-6"], f"形态差异登记：{ids}"
        for diff in block["differences"]:
            for key in ("name", "what", "why_it_matters", "how_to_verify"):
                assert diff.get(key), f"{diff['id']}: 缺 {key}"
                assert len(str(diff[key])) > 20, f"{diff['id']}.{key} 过短"

    def test_blocking_preconditions_are_complete(self, manifest_slice: dict) -> None:
        bps = manifest_slice["blocking_preconditions"]
        ids = [bp["id"] for bp in bps]
        assert ids == [f"BP-{i}" for i in range(1, 11)], f"阻断项编号不连续：{ids}"
        for bp in bps:
            for key in ("blocks", "what", "status", "must_fix_before", "owner_task", "source_refs"):
                assert bp.get(key), f"{bp['id']}: 缺 {key}"
            assert bp.get("consequence") or bp.get("observable_consequences"), (
                f"{bp['id']}: 后果必须写明"
            )
            assert bp["status"] in ("REGISTERED_NOT_FIXED", "PARTIALLY_FIXED_AC_1_4_DONE"), (
                f"{bp['id']}: 未知 status={bp['status']!r}"
            )
            assert bp.get("why_not_fixed_here"), f"{bp['id']}: 缺 why_not_fixed_here"
            for ref in bp["source_refs"]:
                assert _resolve_repo(ref).exists(), f"{bp['id']}: source_ref 不存在 {ref}"


# ════════════════════════════════════════════════════════════════════════════
# Property 70：不得跨 entry 复用
# ════════════════════════════════════════════════════════════════════════════
class TestProperty70NoCrossEntryReuse:
    """**Validates: Requirements 12.10**"""

    def test_cross_entry_isolation_block_present(self, manifest_slice: dict) -> None:
        block = manifest_slice["cross_entry_isolation"]
        assert block["rule"]
        assert isinstance(block["assertions"], list) and len(block["assertions"]) >= 6

    def test_h_entries_disjoint_from_the_four_sibling_slices(self, manifest_slice: dict) -> None:
        mine = {e["entry_id"] for e in manifest_slice["independent_entries"]}
        for path in (D_SLICE_PATH, E_SLICE_PATH, F_SLICE_PATH, G_SLICE_PATH):
            other = {e["entry_id"] for e in _load(path)["independent_entries"]}
            assert not (mine & other), f"与 {path.name} 有交集：{sorted(mine & other)}"
        siblings = set(manifest_slice["sibling_slices"])
        assert siblings == {
            "backend/data/workpaper_sync_d_cycle_manifest_slice.json",
            "backend/data/workpaper_sync_e_cycle_manifest_slice.json",
            "backend/data/workpaper_sync_f_cycle_manifest_slice.json",
            "backend/data/workpaper_sync_g_cycle_manifest_slice.json",
        }
        for rel in siblings:
            assert (ROOT / rel).exists()

    def test_no_entry_carries_a_per_entry_contract_field(self, manifest_slice: dict) -> None:
        for entry in manifest_slice["independent_entries"]:
            blob = json.dumps(entry, ensure_ascii=False)
            assert "contract_id" not in blob, f"{entry['entry_id']}: 出现 contract_id 前向引用"
            assert "definition_bundle_id" not in blob

    def test_deletion_plan_matches_slice(self, manifest_slice: dict, deletion_plan: dict) -> None:
        slice_by_id = {e["entry_id"]: e for e in manifest_slice["independent_entries"]}
        plan_by_id = {e["entry_id"]: e for e in deletion_plan["entries"]}
        assert set(slice_by_id) == set(plan_by_id), "deletion plan 与 slice 的 entry 集合不等"
        for entry_id, plan in plan_by_id.items():
            src = slice_by_id[entry_id]
            assert plan["capability_adjudication"] == src["capability"]
            assert plan["capability_verdict_stage"] == src["capability_verdict_stage"]
            assert plan["capability_target"] == src["capability_target"]
            assert plan["html_counterpart_verdict"] == src["html_counterpart_verdict"]
            assert plan["host_component"] == src["host_path"]
            assert plan["wp_code_pattern"] == src["wp_code_pattern"]
            for ref in plan["html_counterpart_source_refs"]:
                assert _resolve_repo(ref).exists(), f"{entry_id}: plan 的 source_ref 不存在 {ref}"

    def test_deletion_plan_composables_are_distinct_and_real(self, deletion_plan: dict) -> None:
        seen: set[str] = set()
        for entry in deletion_plan["entries"]:
            for item in entry["legacy_composables_to_delete"]:
                path = _resolve_repo(item["file"])
                assert path.exists(), f"{entry['entry_id']}: 待删文件不存在 {item['file']}"
                assert item["file"] not in seen, f"{item['file']} 在 plan 里出现多次"
                seen.add(item["file"])
                assert item["wraps_shared_base"] is False
                actual = len(path.read_text(encoding="utf-8").splitlines())
                assert actual == item["lines"], (
                    f"{item['file']}: 行数声明 {item['lines']} != 实测 {actual}"
                )
        assert len(seen) == 13, (
            f"待删 composable 实测 {len(seen)} 个 —— 9 个 dual-mode + 4 个孤儿 FormData"
        )

    def test_deletion_plan_reasons_are_not_circular(
        self, deletion_plan: dict, paradigm: dict
    ) -> None:
        markers = [
            str(m).lower()
            for ap in paradigm["adjudication_criteria"]["anti_patterns"]
            if ap.get("id") == "AP-1"
            for m in (ap.get("circular_reason_markers") or [])
        ]
        for entry in deletion_plan["entries"]:
            reason = str(entry["adjudication_reason"]).lower()
            hit = [m for m in markers if m in reason]
            if hit:
                assert "html 对端" in reason or "html_counterpart_verdict" in reason, (
                    f"{entry['entry_id']}: plan 的理由命中循环论证标记 {hit[:3]}"
                )

    def test_deletion_plan_declares_no_pilot_and_preserves_shared_base(
        self, deletion_plan: dict
    ) -> None:
        assert deletion_plan["pilot_already_migrated"]["entry_id"] == H1_ENTRY_ID
        assert deletion_plan["deletion_execution_owner"]
        assert deletion_plan["adjudication_source_of_truth"].startswith(
            "backend/data/workpaper_sync_h_cycle_manifest_slice.json"
        )
        for entry in deletion_plan["entries"]:
            for item in entry["legacy_composables_to_delete"]:
                assert "useWorkpaperEntryDualMode" not in item["file"]
        assert len(deletion_plan["excluded_from_plan"]) >= 4

    def test_deletion_plan_child_tab_consumers_are_real(self, deletion_plan: dict) -> None:
        """H4/H8 的子 Tab 消费方必须真实 —— 漏掉即删除时打断子入口。"""
        for entry in deletion_plan["entries"]:
            for item in entry["legacy_composables_to_delete"]:
                stem = pathlib.Path(item["file"]).stem
                production = [
                    c
                    for c in _import_specifier_consumers(stem)
                    if "__tests__" not in c and not c.endswith(".spec.ts")
                ]
                assert sorted(item["consumers"]) == sorted(production), (
                    f"{item['file']}: plan 声明消费方 {item['consumers']} != 实测 {production}"
                )


# ════════════════════════════════════════════════════════════════════════════
# AC 1.4：诚实的模式可见性（step 11）
# ════════════════════════════════════════════════════════════════════════════
class TestAc14HonestModeVisibility:
    """**Validates: Requirements 1.4**

    未注册 adapter 的入口不得显示「可双向回写」，且必须显示**可操作原因**。
    判据落到模板形态（宿主 import + 标签 + entry-id 绑定 + 位于声明的工具栏区块内），
    而不是「只声明 capability」—— 后者是结构性死代码，数据层守卫看不见。
    """

    def test_notice_module_is_the_single_source(self, manifest_slice: dict) -> None:
        source = NOTICE_MODULE.read_text(encoding="utf-8")
        for name in (
            "ENTRY_SYNC_NOTICE_LABEL",
            "ENTRY_SYNC_NOTICE_SUMMARY",
            "ENTRY_SYNC_NOTICE_REASON",
            "SYNC_ADAPTER_REGISTERED_ENTRY_IDS",
            "export function entrySyncNotice",
        ):
            assert name in source, f"通知真源缺 {name}"
        for entry in manifest_slice["independent_entries"]:
            for ref in entry["ui_gate_source_refs"]:
                assert _resolve_repo(ref).exists(), f"{entry['entry_id']}: ui_gate_source_ref 不存在"
            assert any("workpaperEntrySyncNotice.ts" in r for r in entry["ui_gate_source_refs"]), (
                f"{entry['entry_id']}: ui_gate_source_refs 未指向文案真源 ⇒ 可能抄了第二份"
            )

    def test_notice_component_renders_summary_and_binds_entry_id(self) -> None:
        source = NOTICE_COMPONENT.read_text(encoding="utf-8")
        template = _vue_template(source)
        assert "notice.summary" in template, (
            "常显摘要不在模板里 —— 只放 el-tooltip 不算「显示可操作原因」（EP tooltip 是 "
            "teleport 且仅 hover 后进 DOM）"
        )
        assert "notice.reason" in template
        assert ":data-entry-sync-notice=\"props.entryId\"" in template, (
            "缺 data-entry-sync-notice 绑定 ⇒ DOM 级判据无锚点"
        )
        assert "entrySyncNotice" in source

    def test_every_pending_entry_host_mounts_the_notice(self, manifest_slice: dict) -> None:
        for entry in manifest_slice["independent_entries"]:
            if entry["adapter_id"] is not None:
                continue
            host = _resolve_repo(entry["host_path"])
            source = host.read_text(encoding="utf-8")
            assert f"'./sync/{NOTICE_COMPONENT_NAME}.vue'" in source, (
                f"{entry['entry_id']}: 宿主没有 import {NOTICE_COMPONENT_NAME}"
            )
            template = _vue_template(source)
            assert f"<{NOTICE_COMPONENT_NAME}" in template, (
                f"{entry['entry_id']}: 宿主模板里零引用 {NOTICE_COMPONENT_NAME} ⇒ 结构性死代码"
            )
            assert f'entry-id="{entry["entry_id"]}"' in template, (
                f"{entry['entry_id']}: 通知未绑定本 entry 的 entry_id（可能抄了别的 entry）"
            )

    def test_notice_mount_sits_inside_the_declared_mode_toolbar(self, manifest_slice: dict) -> None:
        for entry in manifest_slice["independent_entries"]:
            host = _resolve_repo(entry["host_path"])
            template = _vue_template(host.read_text(encoding="utf-8"))
            block = _toolbar_block(template, entry["ui_toolbar_gate"])
            assert f"<{NOTICE_COMPONENT_NAME}" in block, (
                f"{entry['entry_id']}: 通知不在声明的工具栏区块 "
                f"{entry['ui_toolbar_gate']!r} 内 ⇒ 用户切换模式时看不到原因"
            )
            assert f'entry-id="{entry["entry_id"]}"' in block

    def test_registered_entry_ids_agree_with_the_slice(self, manifest_slice: dict) -> None:
        source = NOTICE_MODULE.read_text(encoding="utf-8")
        match = re.search(
            r"SYNC_ADAPTER_REGISTERED_ENTRY_IDS:\s*readonly string\[\]\s*=\s*\[([\s\S]*?)\]", source
        )
        assert match, "找不到 SYNC_ADAPTER_REGISTERED_ENTRY_IDS 声明"
        registered = set(re.findall(r"['\"]([^'\"]+)['\"]", match.group(1)))
        for entry in manifest_slice["independent_entries"]:
            has_adapter = entry["adapter_id"] is not None
            assert (entry["entry_id"] in registered) == has_adapter, (
                f"{entry['entry_id']}: 通知真源的已注册集合与 slice 的 adapter_id 双口径"
            )

    def test_hosts_do_not_claim_bidirectional_writeback(self, manifest_slice: dict) -> None:
        forbidden = ("可双向回写", "双向同步已启用", "两侧已同步")
        for entry in manifest_slice["independent_entries"]:
            template = _vue_template(_resolve_repo(entry["host_path"]).read_text(encoding="utf-8"))
            for word in forbidden:
                assert word not in template, (
                    f"{entry['entry_id']}: 宿主模板出现 {word!r} —— 未注册 adapter 不得宣称双向"
                )


# ════════════════════════════════════════════════════════════════════════════
# 范式合规
# ════════════════════════════════════════════════════════════════════════════
class TestParadigmCompliance:
    """**Validates: Requirements 12.4**"""

    def test_task50_is_in_both_paradigm_scopes(self, paradigm: dict) -> None:
        assert 50 in paradigm["definition_producer_paradigm"]["applies_to_tasks"]
        assert 50 in paradigm["slice_schema"]["applies_to_tasks"]
        assert len(paradigm["paradigm"]["steps"]) == 7, (
            "冻结的 legacy_deletion_paradigm 步数变了 —— 它是字节冻结的"
        )
        assert len(paradigm["definition_producer_paradigm"]["steps"]) == 12

    def test_capability_enum_comes_from_the_paradigm(self, paradigm: dict) -> None:
        assert tuple(paradigm["adjudication_criteria"]["capability_enum"]) == CAPABILITY_ENUM
        ap1 = next(
            ap for ap in paradigm["adjudication_criteria"]["anti_patterns"] if ap["id"] == "AP-1"
        )
        assert tuple(ap1["allowed_verdict_values"]) == HTML_COUNTERPART_VERDICTS

    def test_single_onlyoffice_illegal_criteria_are_registered(self, paradigm: dict) -> None:
        illegal = paradigm["adjudication_criteria"]["verdicts"]["single_onlyoffice"]["illegal_criteria"]
        for text in ("无 bidirectional adapter", "无 per-entry contract"):
            assert text in illegal

    def test_the_slice_passes_the_paradigm_nominated_validator(self, paradigm: dict) -> None:
        """范式点名的校验器对本 slice 违规数必须为 0，且校验器不能被短路成恒返回空。"""
        import copy
        import importlib.util

        ref = paradigm["slice_schema"]["validator"]
        rel, _, symbol = ref.partition("::")
        module_path = ROOT / rel
        assert module_path.exists(), f"范式点名的校验器文件不存在 {rel}"
        spec = importlib.util.spec_from_file_location("_task50_validator_host", module_path)
        assert spec and spec.loader
        module = importlib.util.module_from_spec(spec)
        spec.loader.exec_module(module)
        validate = getattr(module, symbol)

        slice_doc = _load(MANIFEST_SLICE_PATH)
        problems = validate(slice_doc)
        assert problems == [], (
            f"本 slice 通不过范式自己的 slice_schema（{len(problems)} 条）：\n"
            + "\n".join("  !! " + p for p in problems)
        )

        # 反向自检：抽掉待裁决三字段之一，校验器必须点名它（否则上面那条恒绿）
        mutated = copy.deepcopy(slice_doc)
        mutated["independent_entries"][0].pop("capability_verdict_stage")
        mutated_problems = validate(mutated)
        assert any("capability_verdict_stage" in p for p in mutated_problems), (
            "抽掉 capability_verdict_stage 后校验器没有点名它 ⇒ 校验器已被短路"
        )

    def test_slice_satisfies_the_schema_required_fields(self, paradigm: dict, manifest_slice: dict) -> None:
        schema = paradigm["slice_schema"]
        for key in schema["required_top_level"]:
            assert key in manifest_slice, f"slice 缺顶层必填字段 {key}"
        assert manifest_slice["schema_version"] == schema["target_schema_version"]
        for key in schema["effective_from_task_48"]["required_entry"]:
            for entry in manifest_slice["independent_entries"]:
                assert key in entry, f"{entry['entry_id']}: 缺 Task 48 起必填字段 {key}"

    def test_conditional_sections_that_are_present_satisfy_their_schema(
        self, paradigm: dict, manifest_slice: dict
    ) -> None:
        sections = {c["section"]: c for c in paradigm["slice_schema"]["conditional_sections"]}
        assert "dynamic_row_identity" in manifest_slice, (
            "9 条 entry 全含动态行表，dynamic_row_identity 必须出现"
        )
        assert "parent_duplicate_summary" in manifest_slice, (
            "H4/H8 有 5 条 parent_duplicate 子入口，parent_duplicate_summary 必须出现"
        )
        assert "currency_variant_model" not in manifest_slice, (
            "H3/H7 的双计量模式两套 sheet **各有**独立键，不构成 currency_variant_model 的触发"
            "条件（它要求共用一个持久化键）—— 出现即是伪造声明"
        )
        for name in ("dynamic_row_identity", "parent_duplicate_summary"):
            spec_ = sections[name]
            body = manifest_slice[name]
            for key in spec_["required_fields"]:
                assert key in body, f"{name}: 缺必填字段 {key}"
            for key in spec_.get("required_table_fields") or []:
                for table in body.get("tables") or []:
                    assert key in table, f"{name}.tables: 缺 {key}"
            for key in spec_.get("required_row_identity_fields") or []:
                for table in body.get("tables") or []:
                    assert key in table["row_identity"], f"{name}.row_identity: 缺 {key}"

    def test_paradigm_conflict_is_registered_and_not_stale(self, manifest_slice: dict) -> None:
        block = manifest_slice["paradigm_schema_conflict"]
        assert block["residual_inconsistency"] is None
        assert block["known_red_it_causes"]["status"] == "none"
        assert block["known_red_it_causes"]["locked_by"].startswith(
            _THIS.relative_to(ROOT).as_posix() + "::"
        )
        assert "修改范式 JSON" in block["how_resolved_here"]
        assert "未扩充 capability_enum" in block["how_resolved_here"]

    def test_deletion_plan_has_paradigm_ref(self, deletion_plan: dict) -> None:
        assert deletion_plan["paradigm_ref"] == "backend/data/workpaper_sync_migration_paradigm.json"
        assert (ROOT / deletion_plan["paradigm_ref"]).exists()

    def test_properties_and_requirements_are_declared(
        self, manifest_slice: dict, deletion_plan: dict
    ) -> None:
        assert manifest_slice["properties_verified"] == [22, 23, 69, 70], (
            "Task 50 正文点名 Property 22 / 23 / 69 / 70"
        )
        assert deletion_plan["properties_verified"] == manifest_slice["properties_verified"]
        for req in ("6.4", "6.5", "12.1", "12.4", "12.10", "12.11", "12.12", "14.1"):
            assert req in manifest_slice["requirements_covered"], f"缺 Requirement {req}"
        assert manifest_slice["requirements_covered"] == deletion_plan["requirements_covered"]


# ════════════════════════════════════════════════════════════════════════════
# 源码结构（宿主与 componentType 可达性）
# ════════════════════════════════════════════════════════════════════════════
class TestSourceCodeStructure:
    """**Validates: Requirements 1.4, 12.11**"""

    def test_hosts_exist_and_are_reachable_from_the_renderer_registry(
        self, manifest_slice: dict
    ) -> None:
        registry = HTML_RENDERER_REGISTRY.read_text(encoding="utf-8")
        for entry in manifest_slice["independent_entries"]:
            host = _resolve_repo(entry["host_path"])
            assert host.exists(), f"{entry['entry_id']}: 宿主不存在"
            component = host.stem
            assert re.search(
                rf"const {re.escape(component)} = defineAsyncComponent\(\(\) => import\('\./{re.escape(component)}\.vue'\)\)",
                registry,
            ), f"{entry['entry_id']}: htmlRendererRegistry 里没有 {component} 的 import 边"
            assert re.search(rf"component:\s*{re.escape(component)}\b", registry), (
                f"{entry['entry_id']}: htmlRendererRegistry 里没有 {component} 的 component 绑定"
                " ⇒ 不可达（那就该裁 unreachable 而不是待裁决）"
            )

    def test_hosts_still_import_their_legacy_composable_or_inline_it(
        self, manifest_slice: dict
    ) -> None:
        """删除动作归 Task 66/72 ⇒ 此刻 legacy 必须还在（否则 deletion plan 已过期）。"""
        for entry in manifest_slice["independent_entries"]:
            legacy = entry["legacy_dual_mode"]
            module = _resolve_repo(legacy["module"])
            assert module.exists(), f"{entry['entry_id']}: legacy composable 已被删除，plan 过期"
            # 🔴 必须剥注释：H9 宿主的注释里逐字写着「useH9DualMode composable 后续创建」，
            #    按原文判会把「注释提到」误当成「真 import」（该 entry 的 host_consumes_it=false）。
            host_source = _strip_ts_comments(
                _resolve_repo(entry["host_path"]).read_text(encoding="utf-8")
            )
            if legacy["host_consumes_it"]:
                assert module.stem in host_source, (
                    f"{entry['entry_id']}: 声明宿主消费 legacy，源码里找不到"
                )
            else:
                assert module.stem not in host_source, (
                    f"{entry['entry_id']}: 声明宿主不消费 legacy，源码（剥注释后）里却有 import"
                )
                assert "currentMode" in host_source, (
                    f"{entry['entry_id']}: 宿主既不用 legacy 也没有内联 currentMode ⇒ 无模式切换载体"
                )

    def test_h_cycle_templates_exist(self, manifest_slice: dict) -> None:
        for entry in manifest_slice["independent_entries"]:
            assert (TEMPLATE_DIR / entry["template_ref"]).exists()

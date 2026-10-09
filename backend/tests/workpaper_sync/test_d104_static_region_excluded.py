# -*- coding: utf-8 -*-
"""🔴 T7 裁决 A 的**排除态守卫** —— D1-4 第三区静态 table 已撤回，钉住"它真的不在契约里"。

spec: d1-sync-row-table-engine-and-d1-coverage · T7 裁决 A（跟随 D4 policy check sheet）

═══ 本文件与 O/P/T 三节静态区判据的关系 ═══

O 节 `test_d1_static_region_contract.py` / P 节 `test_d1_static_region_real_stack.py`
是「静态 table **在**契约里」时的判据；T 节 `test_d104_static_region_blocks_row_insertion.py`
证明了「在契约里 ⇒ 动态区不能插行」的回退。
裁决 A 把开关 `_INCLUDE_D104_NOTETYPE_STATIC` 关掉后：
* 那三个文件的**前提消失**（通路不存在），已整体 `skipif(开关 False)` —— 不删，
  marker-relative content 字段落地、开关翻回 `True` 时它们自动复活；
* **本文件反过来守排除态**：开关关时，契约不得含该 table、`all_store_item_ids` 不得含该 item、
  投影/回写清单必须为空、且 D1-4 动态区**恢复可插行**（回退消失）。

═══ 失效条目反向检查（本文件不会烂掉的机制）═══

`test_switch_is_off_and_would_restore_when_flipped` 直接断言开关当前是 `False`。
将来谁把它翻回 `True`（marker-relative 落地）：
* 本文件这条立刻转红 ⇒ 逼迫「要么删本排除守卫、要么把 O/P/T 的 skipif 去掉」，二选一，
  不能两边都留着自相矛盾；
* 同时 O/P/T 的 skipif 失效、那些判据复活 ⇒ 恢复的正确性由它们保证。
两个方向都有判据接住，撤回与恢复都不静默。

🔴 期望值全用**字面量**（不从 spec/常量读被守对象），凡 `for x in <生产常量>` 形态的断言
本仓库反复出现「测试镜像同款错误 ⇒ 恒绿」，不重犯。
"""
from __future__ import annotations

import copy
import hashlib
import io
import json
import sys
import zipfile
from functools import lru_cache
from pathlib import Path
from typing import Any

import openpyxl
import pytest

_BACKEND = Path(__file__).resolve().parents[2]
if str(_BACKEND) not in sys.path:
    sys.path.insert(0, str(_BACKEND))

from app.services.workpaper_sync import excel_instrumentation as EI  # noqa: E402
from app.services.workpaper_sync import excel_materialize as M  # noqa: E402
from app.services.workpaper_sync import phase5_d1_04_bad_debt as D104  # noqa: E402
from app.services.workpaper_sync import phase5_d1_expansion as EXP  # noqa: E402
from app.services.workpaper_sync import phase5_d1_notes_receivable as P  # noqa: E402
from app.services.workpaper_sync.contracts import parse_contract  # noqa: E402
from app.services.workpaper_sync.excel_extract import (  # noqa: E402
    ExcelIdentityBinding,
    _scan_row_identities,
    read_runtime_binding_pairs,
)

SHEET = D104.MANAGED_SHEET_D104
STATIC_TABLE_KEY = "bad_debt_notetype_rows"        # 字面量
STATIC_STORE_ITEM = "D1-bd-notetype-rows"          # 字面量
INDIVIDUAL_ROWS = (13, 16)                         # 动态区容量 4 行
EXPECTED_INSERT_AT = 17
EXPECTED_STYLE_FROM = 16


# ═══════════════════════════════════════════════════════════════════════
# 1. 开关状态 —— 失效条目反向检查
# ═══════════════════════════════════════════════════════════════════════


def test_switch_is_off_and_would_restore_when_flipped() -> None:
    """裁决 A：开关必须为 False。翻回 True（marker-relative 落地）时本条转红。

    转红时的两件事（二选一，不能都留）：
      * 删本排除守卫文件 + 去掉 O/P/T 三节的 `skipif` ⇒ 恢复"在契约里"的判据集；
      * 或维持撤回、把翻转改回来。
    """
    assert EXP._INCLUDE_D104_NOTETYPE_STATIC is False, (
        "_INCLUDE_D104_NOTETYPE_STATIC 已被翻回 True —— 若是 marker-relative content "
        "字段落地后有意恢复，请删除本排除守卫并去掉 O/P/T 三节静态区判据的 skipif；"
        "否则这是误改（撤回是 T7 裁决 A）"
    )


# ═══════════════════════════════════════════════════════════════════════
# 2. 排除态正向守卫 —— 该 table / item / 声明都不在
# ═══════════════════════════════════════════════════════════════════════


def test_contract_has_18_tables_and_no_notetype_static_table() -> None:
    """契约回到 12 sheet / 18 table，且不含 notetype 静态 table。"""
    contract = parse_contract(P.build_contract_payload(), adapter_id=P.ADAPTER_ID)
    sheets = len(contract.sheets)
    table_keys = {t.table_key for s in contract.sheets for t in s.tables}
    assert sheets == 12, f"sheet={sheets}，期望 12"
    assert len(table_keys) == 18, f"table={len(table_keys)}，期望 18（撤回静态 table 后）"
    assert STATIC_TABLE_KEY not in table_keys, (
        f"{STATIC_TABLE_KEY} 仍在契约里 —— 开关关了却没空转，通路门控有漏"
    )


def test_static_item_not_in_all_store_item_ids() -> None:
    """`all_store_item_ids` = 17，不含 notetype item。"""
    items = EXP.all_store_item_ids()
    assert len(items) == 17, f"store item={len(items)}，期望 17"
    assert STATIC_STORE_ITEM not in items, f"{STATIC_STORE_ITEM} 仍在单一口径清单里"


def test_static_region_payloads_and_declarations_are_empty() -> None:
    """契约 payload 侧 + instrumentation 侧的静态声明都空转（同源同开关）。"""
    assert EXP.static_region_table_payloads() == ()
    assert EXP._static_sheet_declarations() == ()


def test_projection_and_merge_registries_are_empty() -> None:
    """投影/回写清单为空 —— 与契约侧成对空转（否则会试图投一张不存在的 table）。"""
    assert P._static_region_projection_builders() == ()
    assert P._static_region_merge_handlers() == ()


def test_combined_projection_skips_notetype_even_if_payload_present() -> None:
    """就算前端仍在写 notetype 载荷，combined projection 也不该产出它的 stable key。

    HTML 侧照常读写该键（走 checklist_responses，与契约无关）—— 这里只验**契约通路**不碰它。
    """
    contract = parse_contract(P.build_contract_payload(), adapter_id=P.ADAPTER_ID)
    payload = json.dumps(
        [{"rowId": "fixed-bank", "noteType": "银行承兑汇票小计", "isFixed": True,
          "priorUnadjusted": 1000, "currentUnadjusted": 2000}],
        ensure_ascii=False,
    )
    proj = P.build_combined_store_projection({STATIC_STORE_ITEM: payload}, contract=contract)
    leaked = [k for k in proj.values if STATIC_TABLE_KEY in k]
    assert leaked == [], f"combined projection 泄漏了 notetype stable key: {leaked[:3]}"


# ═══════════════════════════════════════════════════════════════════════
# 3. 回退消失 —— 撤回换回来的能力：动态区可插行（这是裁决 A 的收益）
# ═══════════════════════════════════════════════════════════════════════


@lru_cache(maxsize=1)
def _substrate() -> bytes:
    return EI.instrument_workbook_bytes_multi(
        P.read_authoritative_template(),
        EXP.instrumentation_specs(),
        gate=P.excel_carrier_gate(),
    ).instrumented_bytes


def _minted_ids(spec) -> list[str]:
    ws = openpyxl.load_workbook(io.BytesIO(_substrate()), data_only=False)[SHEET]
    return [
        str(ws[f"{spec.uuid_col}{r}"].value)
        for r in range(spec.first_data_row, spec.last_data_row + 1)
        if ws[f"{spec.uuid_col}{r}"].value
    ]


def test_dynamic_region_can_insert_row_again_after_withdrawal() -> None:
    """🔴 裁决 A 的**收益**：静态 table 撤回后，D1-4 动态区插行不再被 fail-closed。

    这条与 T 节 `test_experiment_with_static_table_row_insertion_is_refused` 互为镜像 ——
    那条在开关 True 时证明「被拒」，本条在开关 False 时证明「不拒」。
    两条一起把「回退由静态 table 引入」这个因果钉死在两个方向上。
    """
    from app.services.workpaper_sync.excel_extract import _scan_row_identities

    sub = _substrate()
    contract = parse_contract(P.build_contract_payload(), adapter_id=P.ADAPTER_ID)
    spec = D104.SPEC_D104_INDIVIDUAL
    ids = _minted_ids(spec) + ["GTROW-NEW-0001"]  # 超容量 1 行 ⇒ 需要插行
    rows = [
        {"rowId": rid, "label": f"客户{i}", "priorUnadjusted": 100.0 + i,
         "currentProvision": 10.0 + i}
        for i, rid in enumerate(ids, start=1)
    ]
    proj = P.build_combined_store_projection(
        {spec.store_item_id: json.dumps(rows, ensure_ascii=False)}, contract=contract
    )
    binding = ExcelIdentityBinding(
        table_key=spec.table_key, table_name=spec.table_name,
        uuid_column=spec.uuid_col, metadata_sheet="_GT_SYNC",
    )
    ws = openpyxl.load_workbook(io.BytesIO(sub), data_only=False)[SHEET]
    raw_by_row = {
        r: str(ws[f"{spec.uuid_col}{r}"].value or "")
        for r in range(spec.first_data_row, spec.last_data_row + 1)
    }
    table = next(t for s in contract.sheets for t in s.tables if t.table_key == spec.table_key)
    scan = _scan_row_identities(
        table=table, sheet_name=SHEET, uuid_column=spec.uuid_col,
        raw_by_row=raw_by_row, artifact_sha256=hashlib.sha256(sub).hexdigest(),
        tombstoned=(),
    )
    with zipfile.ZipFile(io.BytesIO(sub)) as zf:
        region = M.resolve_managed_region(zf, contract=contract, binding=binding)
    plan = M.plan_managed_writes(
        projection=proj, contract=contract, binding=binding, region=region, scan=scan,
        substrate_entries=M._read_entries(sub),
        substrate_formulas={},
        runtime_binding=read_runtime_binding_pairs(zipfile.ZipFile(io.BytesIO(sub))),
    )
    assert plan.row_shift is not None, "撤回静态 table 后仍算不出插行计划 —— 回退没被解除"
    assert plan.row_shift.insert_at == EXPECTED_INSERT_AT
    assert plan.row_shift.count == 1
    assert plan.row_shift.style_from == EXPECTED_STYLE_FROM


# ═══════════════════════════════════════════════════════════════════════
# 4. HTML 侧不受影响 —— 撤回只动契约通路，前端读写照常
# ═══════════════════════════════════════════════════════════════════════


def test_frontend_store_key_and_readers_are_untouched_by_withdrawal() -> None:
    """撤回是契约层的事；前端那套（load/serialize/read）源文件里那个键必须原样还在。

    这条守的是「裁决 A 没有误伤 HTML 侧」这一前提 —— 它是本裁决成立的依据
    （D1-1 坏账区块的取数只走 HTML store）。
    """
    fe = _BACKEND.parent / "audit-platform/frontend/src/components/workpaper"
    bad_debt = (fe / "composables/useD1BadDebt.ts").read_text(encoding="utf-8")
    model = (fe / "composables/d1AdjudicationModel.ts").read_text(encoding="utf-8")
    # 前端仍以该键读写（不因契约撤回而改）
    assert "D1_BD_NOTETYPE_KEY" in bad_debt
    assert "serializeNoteTypeRows" in bad_debt
    assert "loadNoteTypeRows" in bad_debt
    # D1-1 取数源仍从该键读
    assert "readD1BadDebtByNoteType" in model
    assert "D1_BD_NOTETYPE_KEY = 'D1-bd-notetype-rows'" in model

# -*- coding: utf-8 -*-
"""L3 长期借款真双向接线守卫（spec `l-cycle-true-adapter-registration` · Task 12）。

受管表 `长期借款检查表L3-9`，与 L4 同序：
几何现算 → 契约 → 注册 → HTML 侧对齐 → 真库干净命名空间。

L3 与 L4 的差异（照抄 L4 会错）：
* 受管表**数据区零公式**（公式列为空），L4 有 10 个公式列。
* footer marker 在 **A21**（L4 在 A18），header_rows=2（L4 是三级）。
* 16 字段全标量（L4 是 10 标量 + 14 对数量/金额）。
* 整册 38 definedName / 31 broken + L3-7 四个 #REF!，受管表自身零命中（登记不改）。
"""
from __future__ import annotations

import re

import pytest
from openpyxl import load_workbook

from app.services.workpaper_sync import phase5_l3_long_term_loans as L3
from app.services.workpaper_sync import phase5_l_cycle_common as L
from app.services.workpaper_sync.contracts import parse_contract
from app.services.workpaper_sync.definitions import canonical_digest
from tests.workpaper_sync.l1_adapter_facts import ROOT, _count_bare_if, _is_formula, _query

TEMPLATE = ROOT / "backend" / "wp_templates" / L3.TEMPLATE_RELATIVE_PATH
FRONTEND = ROOT / "audit-platform" / "frontend" / "src" / "components" / "workpaper"
ROWS_TS = FRONTEND / "composables" / "useL3VoucherCheck.ts"


@pytest.fixture(scope="module")
def wb():
    return load_workbook(TEMPLATE, data_only=False)


@pytest.fixture(scope="module")
def ws(wb):
    return wb[L3.MANAGED_SHEET]


# ═══════════════════════════════════════════════════════════════════════════
# 选型与几何（openpyxl 现算）
# ═══════════════════════════════════════════════════════════════════════════


class TestManagedSheetSelection:
    def test_template_sha256_matches_frozen(self) -> None:
        import hashlib

        assert hashlib.sha256(TEMPLATE.read_bytes()).hexdigest() == L3.TEMPLATE_SHA256

    def test_managed_sheet_present_and_rejected_candidates_shape(self, wb) -> None:
        """L3-1 是派生审定表（公式远多于文本）；L3-7 带 #REF!（选型排除理由可复算）。"""
        assert L3.MANAGED_SHEET in wb.sheetnames
        s1 = wb["审定表L3-1"]
        formulas = sum(_is_formula(c.value) for r in s1.iter_rows() for c in r)
        assert formulas >= 80, formulas
        s7 = wb["逾期贷款检查表L3-7"]
        ref_errors = sum(
            isinstance(c.value, str) and "#REF!" in c.value
            for r in s7.iter_rows() for c in r
        )
        assert ref_errors == 4, ref_errors

    def test_geometry(self, ws) -> None:
        assert ws.max_column == 34, "物理 max_column=AH ⇒ UUID 列 AI"
        assert L3.UUID_COL == "AI"
        assert ws[f"A{L3.FOOTER_ROW}"].value == L3.FOOTER_MARKER
        assert (L3.HEADER_GROUP_ROW, L3.HEADER_LEAF_ROW) == (10, 11)
        assert (L3.FIRST_DATA_ROW, L3.LAST_DATA_ROW, L3.FOOTER_ROW) == (12, 20, 21)

    def test_data_region_has_no_formula_columns(self, ws) -> None:
        """L3-9 受管区零公式（与 L4 的 10 公式列相反）。"""
        from openpyxl.utils import get_column_letter

        for r in range(L3.FIRST_DATA_ROW, L3.LAST_DATA_ROW + 1):
            cols = {
                get_column_letter(c)
                for c in range(1, 17)  # A..P 受管面
                if _is_formula(ws.cell(r, c).value)
            }
            assert cols == set(), (r, sorted(cols))
        assert L3.FORMULA_COLUMNS == ()
        assert L3.FORMULA_MASK == ()

    def test_managed_sheet_has_no_bare_if_but_workbook_does(self, wb, ws) -> None:
        """GC-2 per-file 中性化依据：受管表 0、整册 > 0。"""
        assert _count_bare_if(ws) == 0
        assert sum(_count_bare_if(wb[n]) for n in wb.sheetnames) > 0

    def test_sixteen_managed_fields_are_all_editable(self) -> None:
        assert len(L3.MANAGED_FIELD_SPECS_7) == 16
        assert all(mode == "editable" for _, _, mode, *_ in L3.MANAGED_FIELD_SPECS_7)
        cols = [col for _, col, *_ in L3.MANAGED_FIELD_SPECS_7]
        assert cols == list("ABCDEFGHIJKLMNOP")


# ═══════════════════════════════════════════════════════════════════════════
# 契约
# ═══════════════════════════════════════════════════════════════════════════


class TestContract:
    def test_disk_and_source_are_locked(self) -> None:
        contract = L3.assert_contract_file_matches_source()
        assert contract.canonical_sha256 == canonical_digest(L3.build_contract_payload())

    def test_payload_parses_and_has_expected_shape(self) -> None:
        p = L3.build_contract_payload()
        parse_contract(dict(p), adapter_id=L3.ADAPTER_ID)
        assert p["review_status"] == "reviewed"
        assert p["review"]["entry_id"] == L3.ENTRY_ID
        table = p["sheets"][0]["tables"][0]
        assert table["header_rows"] == 2 and table["anchor"] == "A10"
        assert table["footer_anchor"]["search_column"] == "A"
        assert len(table["fields"]) == 16
        assert table["formula_mask"] == []
        assert not any(f["mode"] == "formula" for f in table["fields"])

    def test_derived_readonly_sheet_declared(self) -> None:
        review = L3.build_contract_payload()["review"]
        drs = review["derived_readonly_sheet"]
        assert drs["excel_name"] == L3.MANAGED_SHEET
        assert "R22" in drs["reason"] and "R23" in drs["reason"]

    def test_routes_through_common_skeleton(self) -> None:
        assert L3.EntrySelectionError is L.LEntrySelectionError
        assert L3.build_contract_payload() == L.build_contract_payload(L3.IDENTITY, L3.SPECS)


# ═══════════════════════════════════════════════════════════════════════════
# 注册
# ═══════════════════════════════════════════════════════════════════════════


class TestRegistration:
    def test_provider_aliases_present(self) -> None:
        assert L3.publish_pilot_definitions is L3.publish_definitions
        assert L3.attach_pilot_adapters is L3.attach_adapters
        assert L3.PILOT_WP_CODES == frozenset({"L3L"})

    def test_store_registry_entry(self) -> None:
        from app.services.workpaper_sync.store_item_registry import STORE_MERGE_REGISTRY

        plan = STORE_MERGE_REGISTRY["l3.long_term_loans"]
        assert plan.provider_module == "phase5_l3_long_term_loans"
        assert [i.item_id for i in plan.items] == ["L3-L3-9-voucher-rows"]
        assert plan.oo_crash_neutralization_fn == "neutralize_oo_crash_if_formulas"

    def test_whitelist_and_ledger_are_paired(self) -> None:
        from app.services.workpaper_sync.adapters import registry as R
        from app.services.workpaper_sync.adapters.delivered_contracts_ledger import (
            DELIVERED_PER_ENTRY_CONTRACTS,
        )

        mod = "app.services.workpaper_sync.phase5_l3_long_term_loans"
        assert mod in R._ALLOWED_PROVIDER_MODULES
        rows = [r for r in DELIVERED_PER_ENTRY_CONTRACTS if r["contract_id"] == "l3.long_term_loans"]
        assert len(rows) == 1 and rows[0]["provider_module"] == mod

    def test_wp_code_adjudication_targets_the_whole_workbook_code(self) -> None:
        import json

        doc = json.loads(
            (ROOT / "backend/data/workpaper_sync_entry_wp_code_adjudication.json").read_text("utf-8")
        )
        row = next(a for a in doc["adjudications"] if a["entry_id"] == L3.ENTRY_ID)
        assert row["wp_codes"] == ["L3"] and row["contract_id"] == L3.ADAPTER_ID


# ═══════════════════════════════════════════════════════════════════════════
# HTML 侧对齐（前端 json 键 == 契约 json_pointer 段）
# ═══════════════════════════════════════════════════════════════════════════


class TestHtmlAlignment:
    def test_store_item_id_agrees(self) -> None:
        src = ROWS_TS.read_text("utf-8")
        assert f"L3_VOUCHER_ROWS_ITEM_ID = '{L3.STORE_ITEM_ID}'" in src

    def test_row_identity_uses_shared_minter_not_private_generator(self) -> None:
        """私有 genRowId 已删，改走平台共享 newRowIdentity('l39vc')，已有 rowId 不重铸。"""
        src = ROWS_TS.read_text("utf-8")
        assert "function genRowId" not in src
        assert "newRowIdentity('l39vc')" in src
        assert "raw.rowId || newRowIdentity('l39vc')" in src

    def test_frontend_row_keys_cover_every_contract_field(self) -> None:
        src = ROWS_TS.read_text("utf-8")
        iface = src.split("export interface L3VoucherCheckRow {", 1)[1].split("}", 1)[0]
        frontend = set(re.findall(r"^\s+(\w+)\s*:", iface, re.M))
        contract = {
            f["json_pointer"].rsplit("/", 1)[-1]
            for f in L3.build_contract_payload()["sheets"][0]["tables"][0]["fields"]
        }
        assert contract - frontend == set(), f"契约有、前端缺：{sorted(contract - frontend)}"


# ═══════════════════════════════════════════════════════════════════════════
# 真库：干净命名空间
# ═══════════════════════════════════════════════════════════════════════════


class TestRealDbNamespace:
    def test_new_voucher_key_and_old_positional_key_are_clean(self) -> None:
        rows = _query(
            "SELECT count(*) FILTER (WHERE item_id = 'L3-L3-9-voucher-rows'), "
            "count(*) FILTER (WHERE item_id LIKE 'L3-L3-9-%-data') FROM checklist_responses"
        )
        new, old_positional = rows[0]
        assert old_positional == 0, "L3-9 旧位置化键有真数据 ⇒ 切换前须先迁移"
        assert new >= 0

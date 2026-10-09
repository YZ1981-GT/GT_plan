# -*- coding: utf-8 -*-
"""L1 真双向改线守卫（一）——前序三处失效前提 + 受管表几何基线。

spec: l-cycle-true-adapter-registration
判据: LR-P1 ~ LR-P10

口径真源复用 `tests/workpaper_sync/l1_adapter_facts.py`（唯一真源，禁在本文件复制读数）。
provider / contract / registry 侧判据见 `test_l1_provider_contract_registry.py`。

═══ 本文件守什么 ═══
前序三份 L spec 标 37/37 + 22/23 + 26/27「完成」，但 manifest 里 L 域 8 条 entry 仍全是
`legacy_fake_bidirectional`。本文件把那三处**已失效前提**钉成守卫，防止再次照抄：

1. BP-61-1「三表近空，186 个 planned entry 一个都注册不上」—— 真库实测 12/267/274 行、
   11 个 entry 已跑通 ⇒ 前提失效。🔴 解除理由必须是「entry 覆盖面」而非「表非空」。
2. 「L3/L4/L5 全 NULL、L6/L7/L8 真库 0 行」—— 实测 L2~L8 各 5~8 行非空 remark ⇒ 失效。
3. `l1.short_term_loans.candidate.json` 的 field_mapping —— 字段名与真库键名零交集，
   且把 `审定表L1-1` 的 SUMIF 公式格声明成可写 number ⇒ 照它接线会毁掉整册取数联动。
"""
from __future__ import annotations

from typing import Any

import pytest

from tests.workpaper_sync.l1_adapter_facts import (  # noqa: E402
    BARE_IF_PER_SHEET,
    BARE_IF_TOTAL,
    CANDIDATE_PHANTOM_FIELDS,
    CONTRACT_DIR,
    FOOTER_LABEL_CELL,
    FOOTER_LABEL_TEXT,
    FOOTER_PLACEHOLDER_COLUMNS,
    FOOTER_SUM_COLUMNS,
    L1_ADAPTER_ID,
    L1_ADJ_REAL_FIELDS,
    L1_DERIVED_SHEET,
    L1_ENTRY_ID,
    L1_MANAGED_SHEET,
    MANAGED_FORMULA_COLUMNS,
    MANAGED_FORMULA_SHAPES,
    MANAGED_GEOMETRY,
    LATEST_L_PAYLOAD,
    LATEST_SUPPLY,
    MEASURED_L_PAYLOAD_2026_09_28,
    MEASURED_SUPPLY_2026_09_28,
    SUPPLY_READINGS,
    _count_bare_if,
    _entry,
    _is_formula,
    _load_sheet,
    _manifest_entries,
    _query,
)

@pytest.fixture(scope="module")
def manifest_entries() -> list[dict[str, Any]]:
    return _manifest_entries()


# ═══════════════════════════════════════════════════════════════════════════
# LR-P1 / LR-P2：BP-61-1 前提失效，且解除理由必须是 entry 覆盖面
# ═══════════════════════════════════════════════════════════════════════════


class TestBP611PremiseInvalidated:
    """**Feature: l-cycle-true-adapter-registration, LR-P1 / LR-P2**

    BP-61-1 原文：`working_paper_sync_entry_state` / `working_paper_content_version` /
    `working_paper_content_representation` 三表近空，186 个 planned entry 一个都注册不上。
    """

    def test_three_tables_are_not_nearly_empty(self) -> None:
        """LR-P1：三表行数现算 ≥ 登记值。"""
        rows = _query(
            "SELECT 'entry_state', count(*) FROM working_paper_sync_entry_state "
            "UNION ALL SELECT 'content_version', count(*) FROM working_paper_content_version "
            "UNION ALL SELECT 'representation', count(*) "
            "FROM working_paper_content_representation"
        )
        live = {k: n for k, n in rows}
        # 🔴 与**最新**登记读数比（append-only 时间线，见 l1_adapter_facts.SUPPLY_READINGS）。
        # 回退时不放宽判据，而是按本断言的提示追加新读数条目并写明原因 —— 2026-10-01 即此。
        reg = LATEST_SUPPLY["row_counts"]
        assert live["entry_state"] >= reg["working_paper_sync_entry_state"], (
            f"entry_state 行数 {live['entry_state']} < 登记 "
            f"{reg['working_paper_sync_entry_state']} ⇒ 登记读数须追加新条目说明回退原因"
        )
        assert live["content_version"] >= reg["working_paper_content_version"]
        assert live["representation"] >= reg["working_paper_content_representation"]

    def test_every_regression_in_the_timeline_is_explained(self) -> None:
        """读数时间线里任一次下降，都必须带 `regressed_from` + `regression_cause`。

        没有这条，「追加一条更小的读数」就成了让守卫变绿的后门。
        """
        for prev, cur in zip(SUPPLY_READINGS, SUPPLY_READINGS[1:]):
            dropped = [
                k for k, v in cur["row_counts"].items() if v < prev["row_counts"][k]
            ]
            if not dropped:
                continue
            assert cur.get("regressed_from") == prev["measured_at"], (
                f"{cur['measured_at']} 读数在 {dropped} 上回退却未声明 regressed_from"
            )
            assert cur.get("regression_cause"), (
                f"{cur['measured_at']} 读数回退却未写 regression_cause"
            )

    def test_entry_state_covers_registered_entries(self) -> None:
        """LR-P1：最新登记过的 entry 必须仍在真库（覆盖面只许增长）。"""
        rows = _query(
            "SELECT DISTINCT entry_id FROM working_paper_sync_entry_state ORDER BY entry_id"
        )
        live = {r[0] for r in rows}
        missing = set(LATEST_SUPPLY["entry_state_entry_ids"]) - live
        assert not missing, f"登记过的 entry 在真库消失：{sorted(missing)}"
        assert L1_ENTRY_ID in live, "本 spec task 7b 首发的 L1 representation 不在 entry_state"

    def test_bp611_verdict_follows_the_readings(self) -> None:
        """LR-P2：09-28 读数判「已解除」，10-01 本地回退后同一裁决函数判「未解除」。

        🔴 这正是「解除理由是 entry 覆盖面」的价值：数据一回退，结论跟着翻，不会假绿。
        平台层的「BP-61-1 已解除」结论来自 09-28 的 11 entry 实证，不受本地库回退影响；
        但**本地**已不能复现它，故如实标注。
        """
        first, latest = SUPPLY_READINGS[0], LATEST_SUPPLY
        assert _bp611_released(
            row_total=sum(first["row_counts"].values()),
            entry_coverage=len(first["entry_state_entry_ids"]),
        ) is True
        assert _bp611_released(
            row_total=sum(latest["row_counts"].values()),
            entry_coverage=len(latest["entry_state_entry_ids"]),
        ) is False

    def test_release_verdict_is_driven_by_entry_coverage_not_row_counts(self) -> None:
        """LR-P2：变异 —— 行数非空但 entry 覆盖为 0 时，结论必须翻回「仍阻塞」。

        判据落在**裁决函数的行为**上，不落在「文档里写了已解除」。
        """
        assert _bp611_released(row_total=553, entry_coverage=11) is True
        assert _bp611_released(row_total=553, entry_coverage=0) is False, (
            "行数非空就判解除 ⇒ 把「链路跑过」误当「entry 可注册」"
        )
        assert _bp611_released(row_total=0, entry_coverage=11) is False, (
            "三表全空却有 entry 覆盖是自相矛盾输入，必须判未解除"
        )

    def test_d3_d5_d6_d7_have_representation_but_manifest_still_legacy(self) -> None:
        """LR-P2：剩余差距在治理动作（manifest 重生成 + capability 裁决），不在平台供给。"""
        rows = _query(
            "SELECT DISTINCT entry_id FROM working_paper_sync_entry_state "
            "WHERE entry_id IN ('xlsx/gt-d3-prepaid-accounts','xlsx/gt-d5-receivables-financing',"
            "'xlsx/gt-d6-contract-assets','xlsx/gt-d7-contract-liabilities')"
        )
        have_rep = {r[0] for r in rows}
        # 期望集合取自最新读数（10-01 本地回退后 D3/D5/D6/D7 均不在库）；
        # 只许「最新读数里有的」仍在，不许凭空要求已登记消失的行回来。
        expected = {
            e for e in LATEST_SUPPLY["entry_state_entry_ids"]
            if e in {
                "xlsx/gt-d3-prepaid-accounts", "xlsx/gt-d5-receivables-financing",
                "xlsx/gt-d6-contract-assets", "xlsx/gt-d7-contract-liabilities",
            }
        }
        assert expected <= have_rep, f"D3/D5/D6/D7 中应在库的 {sorted(expected)}，实得 {sorted(have_rep)}"
        # 结论的 manifest 侧不依赖库：四条在 manifest 仍是 legacy（治理动作尚未执行）
        for entry_id in (
            "xlsx/gt-d3-prepaid-accounts", "xlsx/gt-d5-receivables-financing",
            "xlsx/gt-d6-contract-assets", "xlsx/gt-d7-contract-liabilities",
        ):
            assert _entry(entry_id).get("migration_state") == "legacy_fake_bidirectional"
        for entry_id in sorted(have_rep):
            state = _entry(entry_id).get("migration_state")
            assert state == "legacy_fake_bidirectional", (
                f"{entry_id} 的 manifest 状态已变为 {state!r} ⇒ 治理动作已被他人执行，"
                "本判据需按新事实重写而不是放宽"
            )


def _bp611_released(*, row_total: int, entry_coverage: int) -> bool:
    """BP-61-1 解除裁决：**必须**两个条件同时成立，且 entry 覆盖是主判据。"""
    return row_total > 0 and entry_coverage >= 11


# ═══════════════════════════════════════════════════════════════════════════
# LR-P3 / LR-P4：L2~L8 真库载荷非空（前序「L6/L7/L8 为 0」失效）
# ═══════════════════════════════════════════════════════════════════════════

_L_PAYLOAD_SQL = """
SELECT substring(item_id from '^L[0-9]') AS cyc,
       count(*) AS rows,
       count(NULLIF(btrim(coalesce(remark,'')),'')) AS remark_nonblank,
       count(NULLIF(btrim(coalesce(conclusion,'')),'')) AS conclusion_nonblank
FROM checklist_responses WHERE item_id ~ '^L[0-9]'
GROUP BY 1 ORDER BY 1
"""


class TestLDomainRealPayload:
    """**Feature: l-cycle-true-adapter-registration, LR-P3 / LR-P4**"""

    @pytest.fixture(scope="class")
    def live(self) -> dict[str, dict[str, int]]:
        return {
            cyc: {"rows": r, "remark_nonblank": rn, "conclusion_nonblank": cn}
            for cyc, r, rn, cn in _query(_L_PAYLOAD_SQL)
        }

    def test_l2_through_l8_all_have_nonblank_remark(
        self, live: dict[str, dict[str, int]]
    ) -> None:
        """LR-P3：前序 spec 的「L3/L4/L5 全 NULL、L6/L7/L8 真库 0 行」已失效。"""
        for cyc in ("L2", "L3", "L4", "L5", "L6", "L7", "L8"):
            assert cyc in live, f"{cyc} 在真库无任何行 ⇒ 与登记读数矛盾"
            assert live[cyc]["remark_nonblank"] > 0, (
                f"{cyc} 的 remark 非空数为 0 ⇒ 若属实须追加新读数条目，"
                "不得直接改回前序的「0 行」结论"
            )

    def test_live_payload_agrees_with_registered_readings(
        self, live: dict[str, dict[str, int]]
    ) -> None:
        """LR-P3：登记读数与真库双向一致（只允许增长，减少必须显式说明）。"""
        for cyc, reg in LATEST_L_PAYLOAD.items():
            assert cyc in live, f"登记过的 {cyc} 在真库消失"
            assert live[cyc]["rows"] >= reg["rows"], (
                f"{cyc} 行数 {live[cyc]['rows']} < 登记 {reg['rows']}"
            )
            assert live[cyc]["remark_nonblank"] >= reg["remark_nonblank"]

    def test_conclusion_is_structurally_zero_across_l_domain(
        self, live: dict[str, dict[str, int]]
    ) -> None:
        """LR-P4：`conclusion` 全域 0 —— 这条前序结论**仍成立**，沿用。"""
        for cyc, vals in live.items():
            assert vals["conclusion_nonblank"] == 0, (
                f"{cyc} 出现非空 conclusion ⇒ L 域契约需新增 conclusion 映射"
            )

    def test_conclusion_scanner_is_not_vacuous(self) -> None:
        """LR-P4 变异证明：同一口径在**有** conclusion 的域上必须命中非零。

        🔴 结构性零必须配变异证明，否则「扫描器坏了」与「真的是 0」分不开。
        """
        rows = _query(
            "SELECT count(NULLIF(btrim(coalesce(conclusion,'')),'')) "
            "FROM checklist_responses WHERE item_id ~ '^C[0-9]'"
        )
        assert rows[0][0] > 0, (
            "C 域 conclusion 非空数为 0 ⇒ 扫描口径本身有问题，"
            "L 域的「结构性零」结论不可信（C 轮实测 conclusion 10 : remark 8）"
        )


# ═══════════════════════════════════════════════════════════════════════════
# LR-P5 / LR-P6：candidate contract 的两处实质错误
# ═══════════════════════════════════════════════════════════════════════════


class TestCandidateContractDefects:
    """**Feature: l-cycle-true-adapter-registration, LR-P5 / LR-P6**"""

    def test_real_db_field_segments_are_exactly_eight(self) -> None:
        """LR-P5：真库 `L1-adj-*` 字段段恰 8 个。"""
        rows = _query(
            "SELECT DISTINCT substring(item_id from '^L1-adj-[0-9]+-(.*)$') "
            "FROM checklist_responses WHERE item_id LIKE 'L1-adj-%'"
        )
        live = {r[0] for r in rows if r[0]}
        assert live == set(L1_ADJ_REAL_FIELDS), (
            f"真库字段段与登记不符：多 {sorted(live - L1_ADJ_REAL_FIELDS)}，"
            f"少 {sorted(L1_ADJ_REAL_FIELDS - live)}"
        )

    def test_candidate_declared_fields_hit_nothing_in_the_real_db(self) -> None:
        """LR-P5：candidate 声明的 8 个字段名在真库命中 0 ⇒ 字段名不匹配。"""
        for phantom in CANDIDATE_PHANTOM_FIELDS:
            rows = _query(
                "SELECT count(*) FROM checklist_responses "
                f"WHERE item_id LIKE 'L1-adj-%%-{phantom}'"
            )
            assert rows[0][0] == 0, (
                f"candidate 声明的 {phantom!r} 在真库命中 {rows[0][0]} 行 ⇒ "
                "本判据的前提（字段名不匹配）需重算"
            )
        assert not (set(CANDIDATE_PHANTOM_FIELDS) & set(L1_ADJ_REAL_FIELDS)), (
            "candidate 字段集与真库字段集有交集 ⇒ 不是「完全不匹配」，结论须改写"
        )

    def test_derived_sheet_has_zero_writable_cells(self) -> None:
        """LR-P6：`审定表L1-1` 的 R7~R11 × B..L 无一个可输入格 —— 全是公式。"""
        _, ws = _load_sheet(L1_DERIVED_SHEET)
        writable: list[str] = []
        for r in range(7, 12):
            for col in "BCDEFGHIJKL":
                v = ws[f"{col}{r}"].value
                if v is not None and not _is_formula(v):
                    writable.append(f"{col}{r}={v!r}")
        assert writable == [], (
            f"{L1_DERIVED_SHEET} 出现可输入格 {writable} ⇒ canary 改选依据需重算"
        )

    def test_derived_sheet_pulls_from_managed_sheet_via_sumif(self) -> None:
        """LR-P6：B 列是 SUMIF 且引用受管表 ⇒ 写它就毁掉取数联动。"""
        _, ws = _load_sheet(L1_DERIVED_SHEET)
        b7 = str(ws["B7"].value)
        assert "SUMIF" in b7.upper(), f"B7 不是 SUMIF：{b7!r}"
        assert L1_MANAGED_SHEET in b7, f"B7 未引用 {L1_MANAGED_SHEET}：{b7!r}"
        e7 = str(ws["E7"].value)
        assert e7.replace(" ", "") == "=B7+C7+D7", f"E7 形态变了：{e7!r}"
        k7 = str(ws["K7"].value).upper()
        assert "IF(" in k7 and "IFERROR" not in k7, f"K7 不是裸 IF：{k7!r}"

    def test_writing_the_b_column_would_be_flagged(self) -> None:
        """LR-P6 变异：把 B 列列入可写集合必须打红。"""
        _, ws = _load_sheet(L1_DERIVED_SHEET)
        pretend_writable = ["B7", "C7", "D7"]
        clobbered = [ref for ref in pretend_writable if _is_formula(ws[ref].value)]
        assert clobbered == pretend_writable, (
            "把 B/C/D 列当可写格时应当全部识别为「会覆盖公式」，"
            f"实测只识别出 {clobbered} ⇒ 判据空转"
        )


class TestManagedSheetGeometry:
    """**Feature: l-cycle-true-adapter-registration, LR-P7 ~ LR-P10**"""

    def test_geometry_matches_design(self) -> None:
        """LR-P7：八项几何逐项现算等于 design §2.2。"""
        _, ws = _load_sheet(L1_MANAGED_SHEET)
        assert ws.dimensions == MANAGED_GEOMETRY["dims"]
        assert ws.max_row == MANAGED_GEOMETRY["max_row"]
        assert ws.max_column == MANAGED_GEOMETRY["max_col"]
        assert ws[f"A{MANAGED_GEOMETRY['header_group_row']}"].value == "序号"
        assert ws[f"B{MANAGED_GEOMETRY['header_group_row']}"].value == "借款种类"
        assert ws[f"H{MANAGED_GEOMETRY['header_leaf_row']}"].value == "期初余额"
        assert ws[f"A{MANAGED_GEOMETRY['first_data_row']}"].value == 1

    def test_formula_columns_are_exactly_five(self) -> None:
        """LR-P8：公式列恰 K/R/S/T/U，且形态与实测一致。"""
        _, ws = _load_sheet(L1_MANAGED_SHEET)
        first = MANAGED_GEOMETRY["first_data_row"]
        last = MANAGED_GEOMETRY["last_data_row"]
        found: list[str] = []
        for ci in range(1, ws.max_column + 1):
            from openpyxl.utils import get_column_letter

            col = get_column_letter(ci)
            if any(_is_formula(ws[f"{col}{r}"].value) for r in range(first, last + 1)):
                found.append(col)
        assert tuple(found) == MANAGED_FORMULA_COLUMNS, (
            f"数据区公式列实测 {found}，期望 {list(MANAGED_FORMULA_COLUMNS)}"
        )
        for col, shape in MANAGED_FORMULA_SHAPES.items():
            actual = str(ws[f"{col}{first}"].value).replace(" ", "")
            assert actual == shape, f"{col}{first} 形态 {actual!r} != {shape!r}"

    def test_footer_row_is_all_sum_across_fourteen_columns(self) -> None:
        """LR-P8：footer R26 是全 SUM，H..U 连续 14 列，每格形态 `=SUM(X10:X25)`。"""
        _, ws = _load_sheet(L1_MANAGED_SHEET)
        footer = MANAGED_GEOMETRY["footer_row"]
        first = MANAGED_GEOMETRY["first_data_row"]
        last = MANAGED_GEOMETRY["last_data_row"]
        summed = [
            col
            for col in "ABCDEFGHIJKLMNOPQRSTUVWXYZ"
            if _is_formula(ws[f"{col}{footer}"].value)
        ]
        assert tuple(summed) == FOOTER_SUM_COLUMNS, (
            f"footer 公式列实测 {summed}，期望 {list(FOOTER_SUM_COLUMNS)}"
        )
        for col in FOOTER_SUM_COLUMNS:
            actual = str(ws[f"{col}{footer}"].value).replace(" ", "")
            assert actual == f"=SUM({col}{first}:{col}{last})", (
                f"{col}{footer} 不是纵向 SUM：{actual!r} ⇒ footer 语义变了"
            )

    def test_footer_label_is_measured_not_copied_from_d6(self) -> None:
        """LR-P8：🔴 footer 标签现算 —— 实测在 **C 列**，禁抄 d6 的「合   计」。

        这条同时印证「合计标签定位不能限 A 列」（N 轮有一处在 B 列，L1 这处在 C 列）：
        按 A 列找 marker 会得空，从而误判成「无 footer」。
        """
        _, ws = _load_sheet(L1_MANAGED_SHEET)
        footer = MANAGED_GEOMETRY["footer_row"]
        assert ws["A26"].value is None, "A26 有值 ⇒ 标签列结论须重算"
        assert ws[FOOTER_LABEL_CELL].value == FOOTER_LABEL_TEXT, (
            f"{FOOTER_LABEL_CELL} 实测 {ws[FOOTER_LABEL_CELL].value!r}，"
            f"期望 {FOOTER_LABEL_TEXT!r}"
        )
        assert "合   计" != ws[FOOTER_LABEL_CELL].value, (
            "L1 的 footer 标签不是 d6 的 3 半角空格形态 ⇒ marker 不可跨册照抄"
        )
        for col in FOOTER_PLACEHOLDER_COLUMNS:
            assert ws[f"{col}{footer}"].value == "——", (
                f"{col}{footer} 占位符变了 ⇒ footer 形态基线须更新"
            )

    def test_formula_column_scan_must_exclude_header_rows(self) -> None:
        """LR-P8：K 列在 R3/R4（页眉区）也有公式 ⇒ 扫描必须限定数据区。

        不限定会把页眉的 `=底稿目录!...` 算进数据区公式列，K 列的 `formula_mask`
        跨度就会从 R10 一路错扩到 R3。
        """
        _, ws = _load_sheet(L1_MANAGED_SHEET)
        all_k = [
            r
            for r in range(1, ws.max_row + 1)
            if _is_formula(ws[f"K{r}"].value)
        ]
        assert 3 in all_k and 4 in all_k, "K3/K4 页眉公式消失 ⇒ 本判据前提须重算"
        in_data = [r for r in all_k if 10 <= r <= 25]
        assert len(in_data) == 16, f"K 列数据区公式行数 {len(in_data)} != 16"

    def test_managed_sheet_has_no_bare_if(self) -> None:
        """LR-P9：受管表裸 IF == 0（整册 112 格全在别的 sheet）。"""
        _, ws = _load_sheet(L1_MANAGED_SHEET)
        assert _count_bare_if(ws) == 0, "受管表出现裸 IF ⇒ 中性化策略需重新评估"

    def test_workbook_bare_if_distribution(self) -> None:
        """LR-P9：整册裸 IF 分布逐 sheet 等于 design §4.3。"""
        wb, _ = _load_sheet(L1_MANAGED_SHEET)
        live = {n: _count_bare_if(wb[n]) for n in wb.sheetnames}
        total = sum(live.values())
        assert total == BARE_IF_TOTAL, f"整册裸 IF 实测 {total}，登记 {BARE_IF_TOTAL}"
        for sheet, expected in BARE_IF_PER_SHEET.items():
            assert live.get(sheet) == expected, (
                f"{sheet!r} 裸 IF 实测 {live.get(sheet)}，登记 {expected}"
            )

    def test_defined_names_are_empty_with_mutation_proof(self) -> None:
        """LR-P10：definedName 0 / broken 0，并证明扫描器非空转。"""
        wb, _ = _load_sheet(L1_MANAGED_SHEET)
        names = list(wb.defined_names.keys())
        assert names == [], f"L1 册出现 definedName {names} ⇒ 基线须更新"

        from openpyxl.workbook.defined_name import DefinedName

        wb.defined_names.add(DefinedName("gt_probe", attr_text="#REF!"))
        probed = list(wb.defined_names.keys())
        broken = [k for k in probed if "#REF!" in str(wb.defined_names[k].value)]
        assert probed == ["gt_probe"] and broken == ["gt_probe"], (
            "注入含 #REF! 的 defined name 后扫描器未命中 ⇒ 空分母不可信"
        )



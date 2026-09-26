# -*- coding: utf-8 -*-
"""判据：D1-8/D1-16 四区声明几何正确 + 灰度零回归 + 自动发现进位移判据清单。

spec: d1-sync-row-table-engine-and-d1-coverage · Task 27 · Requirements 5.1 / 5.4
"""
from __future__ import annotations

import io
import os
import sys
import warnings
from pathlib import Path
from typing import Any

import pytest

_REPO = Path(__file__).resolve().parents[3]
_BACKEND = _REPO / "backend"
if str(_BACKEND) not in sys.path:  # pragma: no cover
    sys.path.insert(0, str(_BACKEND))
os.environ.setdefault("DB_DISABLE_SSL", "True")
warnings.filterwarnings("ignore")

from app.services.workpaper_sync import phase5_d1_08_endorsement as D108  # noqa: E402
from app.services.workpaper_sync import phase5_d1_16_writeoff as D116  # noqa: E402
from app.services.workpaper_sync import phase5_d1_expansion as D1E  # noqa: E402
from app.services.workpaper_sync import phase5_d1_notes_receivable as D1  # noqa: E402
from app.services.workpaper_sync.excel_extract import BindingKind  # noqa: E402
from app.services.workpaper_sync.phase5_row_table_sheet import StoreKind  # noqa: E402


@pytest.fixture(scope="module")
def workbook() -> Any:
    import openpyxl
    return openpyxl.load_workbook(io.BytesIO(D1.read_authoritative_template()), data_only=False)


# ═══════════════════════════════════════════════════════════════════════════
# D1-8 贴现/背书 双区
# ═══════════════════════════════════════════════════════════════════════════

class TestD108DiscountSpec:
    """区一：贴现 R14-R21。"""

    def test_managed_sheet_matches_real_tab_name(self, workbook: Any) -> None:
        assert D108.MANAGED_SHEET_D108 in workbook.sheetnames

    def test_geometry(self) -> None:
        s = D108.SPEC_D108_DISCOUNT
        assert s.first_data_row == 14
        assert s.last_data_row == 21
        assert s.footer_row == 22
        assert s.uuid_col == "Q"
        assert s.row_identity_key == "rowId"
        assert s.store_kind is StoreKind.rows

    def test_field_count_matches_template_columns(self) -> None:
        assert len(D108.SPEC_D108_DISCOUNT.field_specs) == 16  # A-P

    def test_store_key_matches_frontend(self) -> None:
        assert D108.SPEC_D108_DISCOUNT.store_item_id == "D1-endorse-discount-rows"

    def test_formula_mask_covers_footer_sum_columns(self) -> None:
        mask = D108.SPEC_D108_DISCOUNT.formula_mask
        assert len(mask) == 4  # E/F/L/M
        for col in ("E", "F", "L", "M"):
            assert any(col in m for m in mask), f"formula_mask 应覆盖 {col} 列"

    def test_footer_marker_codepoints(self) -> None:
        assert D108.FOOTER_MARKER_D108 == "合计"
        assert [hex(ord(c)) for c in D108.FOOTER_MARKER_D108] == ["0x5408", "0x8ba1"]

    def test_footer_marker_matches_template(self, workbook: Any) -> None:
        ws = workbook[D108.MANAGED_SHEET_D108]
        assert ws.cell(row=22, column=1).value == D108.FOOTER_MARKER_D108

    def test_header_text_matches_template(self, workbook: Any) -> None:
        ws = workbook[D108.MANAGED_SHEET_D108]
        for _key, col, _mode, _vt, _jk, label, _gh in D108.SPEC_D108_DISCOUNT.field_specs:
            from openpyxl.utils import column_index_from_string
            ci = column_index_from_string(col)
            real = ws.cell(row=12, column=ci).value
            assert real is not None, f"R12/{col} 应有表头文字，实得 None（字段 {_key}）"


class TestD108TransferSpec:
    """区二：背书 R26-R33。"""

    def test_geometry(self) -> None:
        s = D108.SPEC_D108_TRANSFER
        assert s.first_data_row == 26
        assert s.last_data_row == 33
        assert s.footer_row == 34
        assert s.uuid_col == "R"
        assert s.row_identity_key == "rowId"

    def test_field_count_matches_template_columns(self) -> None:
        assert len(D108.SPEC_D108_TRANSFER.field_specs) == 16  # A-P

    def test_store_key_matches_frontend(self) -> None:
        assert D108.SPEC_D108_TRANSFER.store_item_id == "D1-endorse-transfer-rows"

    def test_formula_mask_covers_footer_sum_columns(self) -> None:
        mask = D108.SPEC_D108_TRANSFER.formula_mask
        assert len(mask) == 1  # E only
        assert "E" in mask[0]

    def test_shared_sheet_key(self) -> None:
        """同 sheet 双区必须共享 sheet_key（同 D1-4 先例）。"""
        assert D108.SPEC_D108_DISCOUNT.sheet_key == D108.SPEC_D108_TRANSFER.sheet_key

    def test_different_uuid_cols(self) -> None:
        """同 sheet 双区必须各用不同 UUID 列（D4-1 教训）。"""
        assert D108.SPEC_D108_DISCOUNT.uuid_col != D108.SPEC_D108_TRANSFER.uuid_col

    def test_different_table_names(self) -> None:
        assert D108.SPEC_D108_DISCOUNT.table_name != D108.SPEC_D108_TRANSFER.table_name

    def test_k_l_m_columns_differ_between_regions(self) -> None:
        """K/L/M 三列两区语义不同（贴现银行 vs 背书转让单位）——column_key 必须不同。"""
        disc_keys = {s[0] for s in D108.SPEC_D108_DISCOUNT.field_specs if s[1] in ("K", "L", "M")}
        xfer_keys = {s[0] for s in D108.SPEC_D108_TRANSFER.field_specs if s[1] in ("K", "L", "M")}
        assert disc_keys.isdisjoint(xfer_keys), (
            f"K/L/M 的 column_key 在两区间应互不相同：贴现={disc_keys} 背书={xfer_keys}"
        )


# ═══════════════════════════════════════════════════════════════════════════
# D1-16 转回/核销 双区
# ═══════════════════════════════════════════════════════════════════════════

class TestD116ReversalSpec:
    """区一：转回 R12-R14。"""

    def test_managed_sheet_matches_real_tab_name(self, workbook: Any) -> None:
        assert D116.MANAGED_SHEET_D116 in workbook.sheetnames

    def test_geometry(self) -> None:
        s = D116.SPEC_D116_REVERSAL
        assert s.first_data_row == 12
        assert s.last_data_row == 14
        assert s.footer_row == 15
        assert s.uuid_col == "I"
        assert s.row_identity_key == "id"  # 🔴 非 rowId
        assert s.store_kind is StoreKind.rows

    def test_field_count_matches_template_columns(self) -> None:
        assert len(D116.SPEC_D116_REVERSAL.field_specs) == 8  # A-H

    def test_store_key_matches_frontend(self) -> None:
        assert D116.SPEC_D116_REVERSAL.store_item_id == "D1-writeoff-reversal-rows"

    def test_formula_mask_covers_footer_sum_columns(self) -> None:
        mask = D116.SPEC_D116_REVERSAL.formula_mask
        assert len(mask) == 2  # E/F
        for col in ("E", "F"):
            assert any(col in m for m in mask)

    def test_row_identity_key_is_id_not_rowId(self) -> None:
        """D1-16 前端用 `id` 而非 `rowId`——声明必须照前端逐字匹配。"""
        assert D116.SPEC_D116_REVERSAL.row_identity_key == "id"
        assert D116.SPEC_D116_WRITEOFF.row_identity_key == "id"


class TestD116WriteoffSpec:
    """区二：核销 R18-R20。"""

    def test_geometry(self) -> None:
        s = D116.SPEC_D116_WRITEOFF
        assert s.first_data_row == 18
        assert s.last_data_row == 20
        assert s.footer_row == 21
        assert s.uuid_col == "J"

    def test_field_count_matches_template_columns(self) -> None:
        assert len(D116.SPEC_D116_WRITEOFF.field_specs) == 8  # A-H

    def test_store_key_matches_frontend(self) -> None:
        assert D116.SPEC_D116_WRITEOFF.store_item_id == "D1-writeoff-writeoff-rows"

    def test_formula_mask_covers_footer_sum_columns(self) -> None:
        mask = D116.SPEC_D116_WRITEOFF.formula_mask
        assert len(mask) == 1  # C only
        assert "C" in mask[0]

    def test_shared_sheet_key(self) -> None:
        assert D116.SPEC_D116_REVERSAL.sheet_key == D116.SPEC_D116_WRITEOFF.sheet_key

    def test_different_uuid_cols(self) -> None:
        assert D116.SPEC_D116_REVERSAL.uuid_col != D116.SPEC_D116_WRITEOFF.uuid_col


# ═══════════════════════════════════════════════════════════════════════════
# 灰度开关零回归
# ═══════════════════════════════════════════════════════════════════════════

class TestExpansionGrayScaleZeroRegression:
    """开关 off ⇒ 与改动前逐字等价（只有 D1-3）。"""

    def test_switches_are_off_by_default(self) -> None:
        assert D1E._INCLUDE_D108_ENDORSEMENT is False
        assert D1E._INCLUDE_D116_WRITEOFF is False

    def test_instrumentation_specs_unchanged_with_switches_off(self) -> None:
        specs = D1E.instrumentation_specs()
        assert len(specs) == 1
        assert str(specs[0].managed_sheet) == "原值明细表（按客户）D1-3"

    def test_store_item_ids_unchanged_with_switches_off(self) -> None:
        items = D1E.all_store_item_ids()
        assert items == ("D1-cust-rows",)


class TestExpansionWithD108D116On:
    """开关 on ⇒ 4 个新 spec 自动出现且几何正确。"""

    def test_instrumentation_grows_to_5_with_both_on(self) -> None:
        original_d108 = D1E._INCLUDE_D108_ENDORSEMENT
        original_d116 = D1E._INCLUDE_D116_WRITEOFF
        try:
            D1E._INCLUDE_D108_ENDORSEMENT = True
            D1E._INCLUDE_D116_WRITEOFF = True
            specs = D1E.instrumentation_specs()
            assert len(specs) == 5
            sheets = [str(s.managed_sheet) for s in specs]
            assert sheets.count(D108.MANAGED_SHEET_D108) == 2  # 贴现+背书
            assert sheets.count(D116.MANAGED_SHEET_D116) == 2  # 转回+核销
        finally:
            D1E._INCLUDE_D108_ENDORSEMENT = original_d108
            D1E._INCLUDE_D116_WRITEOFF = original_d116

    def test_store_items_grow_to_5_with_both_on(self) -> None:
        original_d108 = D1E._INCLUDE_D108_ENDORSEMENT
        original_d116 = D1E._INCLUDE_D116_WRITEOFF
        try:
            D1E._INCLUDE_D108_ENDORSEMENT = True
            D1E._INCLUDE_D116_WRITEOFF = True
            items = D1E.all_store_item_ids()
            assert len(items) == 5
            assert "D1-endorse-discount-rows" in items
            assert "D1-endorse-transfer-rows" in items
            assert "D1-writeoff-reversal-rows" in items
            assert "D1-writeoff-writeoff-rows" in items
        finally:
            D1E._INCLUDE_D108_ENDORSEMENT = original_d108
            D1E._INCLUDE_D116_WRITEOFF = original_d116

    def test_d1_8_auto_enters_multi_region_parametrized_coverage(self) -> None:
        """Task 24 参数化判据的变异反证：D1-8 开关翻转后自动进入 `_multi_region_sheets()`。"""
        original = D1E._INCLUDE_D108_ENDORSEMENT
        try:
            D1E._INCLUDE_D108_ENDORSEMENT = True
            # 复用 Task 24 写的 _multi_region_sheets 逻辑
            sys.path.insert(0, str(_BACKEND / "scripts" / "check"))
            try:
                from check_sync_provider_golden_digest import PROVIDERS  # noqa: F401
            finally:
                sys.path.remove(str(_BACKEND / "scripts" / "check"))
            from tests.workpaper_sync.test_sibling_table_ref_row_shift import (
                _multi_region_sheets,
            )
            multi = _multi_region_sheets()
            assert D108.MANAGED_SHEET_D108 in multi, (
                f"D1-8 开关翻转后应自动进入多区覆盖清单，实得 {sorted(multi)}"
            )
            assert len(multi[D108.MANAGED_SHEET_D108]) == 2
        finally:
            D1E._INCLUDE_D108_ENDORSEMENT = original

    def test_d1_16_auto_enters_multi_region_parametrized_coverage(self) -> None:
        """同 D1-8：D1-16 开关翻转后也应自动进入多区覆盖清单。"""
        original = D1E._INCLUDE_D116_WRITEOFF
        try:
            D1E._INCLUDE_D116_WRITEOFF = True
            sys.path.insert(0, str(_BACKEND / "scripts" / "check"))
            try:
                from check_sync_provider_golden_digest import PROVIDERS  # noqa: F401
            finally:
                sys.path.remove(str(_BACKEND / "scripts" / "check"))
            from tests.workpaper_sync.test_sibling_table_ref_row_shift import (
                _multi_region_sheets,
            )
            multi = _multi_region_sheets()
            assert D116.MANAGED_SHEET_D116 in multi, (
                f"D1-16 开关翻转后应自动进入多区覆盖清单，实得 {sorted(multi)}"
            )
            assert len(multi[D116.MANAGED_SHEET_D116]) == 2
        finally:
            D1E._INCLUDE_D116_WRITEOFF = original

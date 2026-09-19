"""Small real XLSX carrier for publish-hash tests (no database)."""
import io
import uuid
from openpyxl import Workbook
from openpyxl.worksheet.table import Table
from app.services.workpaper_sync.excel_instrumentation import GT_SYNC_SHEET_NAME


def workbook_fixture(sheet_key="g7-disclosure", rows=2):
    anchors = dict(sheet_key=sheet_key, table_name="GT_G1_ROWS",
                   uuid_column_letter="N", metadata_sheet=GT_SYNC_SHEET_NAME)
    wb = Workbook()
    ws = wb.active
    ws.title = "Managed"
    for col in range(1, 15):
        ws.cell(7, col, f"column_{col}")
    for row in range(8, 8 + rows):
        ws.cell(row, 7, f"row_{row}")
        ws.cell(row, 8, row * 10)
        ws.cell(row, 14, str(uuid.UUID(int=row)))
    ws.column_dimensions["N"].hidden = True
    ws.add_table(Table(displayName=anchors["table_name"], ref=f"A7:N{7 + rows}"))
    wb.create_sheet(GT_SYNC_SHEET_NAME).sheet_state = "veryHidden"
    output = io.BytesIO()
    wb.save(output)
    return output.getvalue(), anchors


def attach_projection_json(base_bytes, *, contract_id, values):
    """把 projection JSON 注入 `workbook_fixture()` 产出的受管 xlsx 字节。

    `_StructuredFakeExcelAdapter` 等 harness 需要「带真实受管结构（Excel Table +
    UUID 列 + _GT_SYNC sheet）」且「extract 能真读回投影」的字节：前者保证
    `compute_structure_hash_from_artifact` 反读得出结构，后者保证 roundtrip 等值。
    这里在 fixture zip 里追加 `_gt_sync/projection.json`，两个诉求一次满足。
    """
    import hashlib
    import json
    import zipfile

    buf = io.BytesIO()
    with zipfile.ZipFile(io.BytesIO(base_bytes)) as src, zipfile.ZipFile(
        buf, "w", compression=zipfile.ZIP_DEFLATED
    ) as dst:
        for item in src.infolist():
            dst.writestr(item, src.read(item.filename))
        dst.writestr(
            "_gt_sync/projection.json",
            json.dumps(
                {"contract_id": contract_id, "values": values},
                sort_keys=True,
                ensure_ascii=False,
            ).encode("utf-8"),
        )
    blob = buf.getvalue()
    return blob, hashlib.sha256(blob).hexdigest()


def make_structured_excel_adapter(
    base_cls, *, plain_fn, d_fn, materialize_result_cls, anchors
):
    """工厂：基于 `base_cls`（测试文件里的 `_FakeExcelAdapter`）派生一个写盘带**真实
    受管结构**（Excel Table + UUID 列 + `_GT_SYNC` sheet）的替身。

    父类的 `_minimal_ooxml` 是「能过 Task 11 安全门」的空壳，没有受管 Table，
    `_projection_structure_hash` 里 `compute_structure_hash_from_artifact` 反读不出结构会抛。
    这里复用 `workbook_fixture()`（G1-1 已证它 + `xc` 契约能干净通过
    `assert_no_structure_drift`）的字节，再把 projection JSON 注进同一个 zip —— extract
    仍能真读回投影、roundtrip 等值照跑，而结构哈希也能真算。

    以工厂（传入 `base_cls`）而非 import 派生，避免与测试文件形成循环 import。
    plan 必须传 `workbook_fixture()` 给出的同一组 `structure_anchors`，否则
    `identity_inventory` 反读为空即抛。
    """

    class _StructuredFakeExcelAdapter(base_cls):
        def materialize(self, *, substrate, projection, output, contract):
            self.materialize_calls += 1
            payload = {}
            for key, value in projection.values.items():
                if key in self.drop_keys:
                    continue
                payload[key] = {
                    "value": self.mutate.get(key, plain_fn(value.value)),
                    "value_type": value.value_type.value,
                    "mode": value.mode.value,
                    "row_key": value.row_key,
                }
            for key, value in self.extra_values.items():
                payload[key] = value
            base, _anchors = workbook_fixture(sheet_key=anchors["sheet_key"])
            blob, digest = attach_projection_json(
                base, contract_id=projection.contract_id, values=payload
            )
            output.write_bytes(blob)
            return materialize_result_cls(
                output_path=output,
                document_type=projection.document_type,
                artifact_sha256=digest,
                structure_hash=d_fn("structure"),
                identity_inventory_sha256=d_fn("identity"),
                managed_field_count=len(payload),
            )

    return _StructuredFakeExcelAdapter

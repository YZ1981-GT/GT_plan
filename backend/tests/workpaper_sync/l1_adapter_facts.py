# -*- coding: utf-8 -*-
"""L1 真双向改线守卫的**共享事实层**（口径真源）。

spec: l-cycle-true-adapter-registration

按 `k_foundation_facts.py` / `k_lane1_facts.py` 的既有范式从
`test_l1_adapter_registration.py` 拆出 —— 拆分动因是硬约束：`check_file_size.py`
对**新增** `.py` 强制 ≤800 行（whitelist 只许历史大文件），单文件时为 1039 行。

消费方两个：`test_l1_adapter_registration.py`（前提复核 + 模板几何）与
`test_l1_provider_contract_registry.py`（provider / contract / registry）。
🔴 **实测读数与几何基线只在本模块出现一次** —— 两个测试文件都从这里取，
不各自复制，否则会漂移。
"""
from __future__ import annotations

import json
import os
import warnings
from pathlib import Path
from typing import Any

import pytest

warnings.filterwarnings("ignore", category=UserWarning, module="openpyxl")

ROOT = Path(__file__).resolve().parents[3]
BACKEND = ROOT / "backend"
DATA = BACKEND / "data"
MANIFEST_PATH = DATA / "workpaper_sync_entry_manifest.json"
CONTRACT_DIR = DATA / "workpaper_sync_contracts"
L1_TEMPLATE = BACKEND / "wp_templates" / "L" / "L1 短期借款.xlsx"

L1_ENTRY_ID = "xlsx/gt-l1-short-term-loans"
L1_ADAPTER_ID = "l1.short_term_loans"
L1_MANAGED_SHEET = "明细表L1-2"
L1_DERIVED_SHEET = "审定表L1-1"

# ═══════════════════════════════════════════════════════════════════════════
# 实测读数登记（append-only —— 新读数往后追加，禁改历史条目）
# ═══════════════════════════════════════════════════════════════════════════

#: BP-61-1 三表读数。🔴 只登记行数与 entry 覆盖；「BP-61-1 是否解除」由 LR-P2 裁决。
MEASURED_SUPPLY_2026_09_28: dict[str, Any] = {
    "measured_at": "2026-09-28",
    "row_counts": {
        "working_paper_sync_entry_state": 12,
        "working_paper_content_version": 267,
        "working_paper_content_representation": 274,
    },
    "entry_state_entry_ids": (
        "opaque-017624e2-cd7e-4641-bd99-db7c1f1f5f7e",
        "xlsx/b60/gt-b60-bundle",
        "xlsx/gt-d1-notes-receivable",
        "xlsx/gt-d2-accounts-receivable",
        "xlsx/gt-d3-prepaid-accounts",
        "xlsx/gt-d4-operating-revenue",
        "xlsx/gt-d5-receivables-financing",
        "xlsx/gt-d6-contract-assets",
        "xlsx/gt-d7-contract-liabilities",
        "xlsx/gt-g7-long-term-equity-main",
        "xlsx/gt-h1-fixed-assets",
    ),
}

#: L 域真库载荷。NULL 与空串分开计（铁律⑧）。
MEASURED_L_PAYLOAD_2026_09_28: dict[str, dict[str, int]] = {
    "L0": {"rows": 1, "remark_nonblank": 1, "conclusion_nonblank": 0},
    "L1": {"rows": 33, "remark_nonblank": 33, "conclusion_nonblank": 0},
    "L2": {"rows": 8, "remark_nonblank": 8, "conclusion_nonblank": 0},
    "L3": {"rows": 11, "remark_nonblank": 7, "conclusion_nonblank": 0},
    "L4": {"rows": 6, "remark_nonblank": 6, "conclusion_nonblank": 0},
    "L5": {"rows": 7, "remark_nonblank": 7, "conclusion_nonblank": 0},
    "L6": {"rows": 5, "remark_nonblank": 5, "conclusion_nonblank": 0},
    "L7": {"rows": 5, "remark_nonblank": 5, "conclusion_nonblank": 0},
    "L8": {"rows": 5, "remark_nonblank": 5, "conclusion_nonblank": 0},
}

#: 真库 `L1-adj-*` 的字段段全集（4 分类 × 8 字段 + conclusion 1 = 33 行）。
L1_ADJ_REAL_FIELDS: frozenset[str] = frozenset(
    {
        "aje",
        "audited",
        "beginning",
        "creditAmount",
        "debitAmount",
        "endBalance",
        "rje",
        "unadjusted",
    }
)

#: 🔴 2026-10-01 追加读数：本地 PG **出现非整库的数据回退**（如实登记，成因未查明）。
#:
#: 现象（现算）：`projects` 139 / `working_paper` 997 / `trial_balance` 730 仍在
#: ⇒ 不是重建库；但 `working_paper_sync_entry_state` / `content_version` / `representation`
#: 三表在本轮动手前**全为 0**（09-28 登记 12 / 267 / 274），`checklist_responses` 只剩 42 行、
#: `item_id ~ '^L[0-9]'` **0 行**，而 `definition_artifact` 仍有 36 行（9 entry × 4）。
#: ⇒ 有人（或某个测试夹具）对这几张表做过清理。PG 容器自 2026-05-10 起未重建、卷未换。
#: 🔴 **不推断成因**（测试 TRUNCATE / 并发会话手工清理 / 恢复备份都可能），只登记现象。
#:
#: 本条读数之后的三表值，是本 spec task 7b 在空表上**首发 L1** 的结果（1 / 1 / 1）；
#: L 域载荷是 `backend/scripts/e2e/seed_l_cycle_canary_rows.py --apply` 复现的夹具行
#: （`wp_ref='l-cycle-canary-e2e'`）。L2 比 09-28 少 1 行：那一行是 LC-22 的 G8 污染样本，
#: 夹具**刻意不造**（造出来就是伪造缺陷证据），见该脚本 docstring。
MEASURED_SUPPLY_2026_10_01: dict[str, Any] = {
    "measured_at": "2026-10-01",
    "row_counts": {
        "working_paper_sync_entry_state": 1,
        "working_paper_content_version": 1,
        "working_paper_content_representation": 1,
    },
    "entry_state_entry_ids": ("xlsx/gt-l1-short-term-loans",),
    "regressed_from": "2026-09-28",
    "regression_cause": "unknown_partial_table_cleanup",
}

MEASURED_L_PAYLOAD_2026_10_01: dict[str, dict[str, int]] = {
    "L0": {"rows": 1, "remark_nonblank": 1, "conclusion_nonblank": 0},
    "L1": {"rows": 33, "remark_nonblank": 33, "conclusion_nonblank": 0},
    "L2": {"rows": 7, "remark_nonblank": 7, "conclusion_nonblank": 0},
    "L3": {"rows": 11, "remark_nonblank": 7, "conclusion_nonblank": 0},
    "L4": {"rows": 6, "remark_nonblank": 6, "conclusion_nonblank": 0},
    "L5": {"rows": 7, "remark_nonblank": 7, "conclusion_nonblank": 0},
    "L6": {"rows": 5, "remark_nonblank": 5, "conclusion_nonblank": 0},
    "L7": {"rows": 5, "remark_nonblank": 5, "conclusion_nonblank": 0},
    "L8": {"rows": 5, "remark_nonblank": 5, "conclusion_nonblank": 0},
}

#: 读数时间线（append-only）。守卫一律与**最新**一条比；历史条目只作审计轨迹。
SUPPLY_READINGS: tuple[dict[str, Any], ...] = (
    MEASURED_SUPPLY_2026_09_28,
    MEASURED_SUPPLY_2026_10_01,
)
L_PAYLOAD_READINGS: tuple[dict[str, dict[str, int]], ...] = (
    MEASURED_L_PAYLOAD_2026_09_28,
    MEASURED_L_PAYLOAD_2026_10_01,
)
LATEST_SUPPLY: dict[str, Any] = SUPPLY_READINGS[-1]
LATEST_L_PAYLOAD: dict[str, dict[str, int]] = L_PAYLOAD_READINGS[-1]

#: candidate contract 声明的字段名 —— 真库命中 0，是 LR-P5 的反例集合。
CANDIDATE_PHANTOM_FIELDS: tuple[str, ...] = (
    "beginUnadjusted",
    "beginAje",
    "beginRje",
    "beginAudited",
    "endUnadjusted",
    "endAje",
    "endRje",
    "endAudited",
)


# ═══════════════════════════════════════════════════════════════════════════
# helper：真库直连 / manifest / 模板
# ═══════════════════════════════════════════════════════════════════════════


def _dsn() -> str:
    """把 app 配置里的 asyncpg URL 转成 psycopg2 可用的 DSN。"""
    url = os.environ.get("DATABASE_URL", "")
    if not url:
        try:
            from app.core.config import settings  # type: ignore[attr-defined]

            url = str(settings.DATABASE_URL)
        except Exception:  # pragma: no cover - 配置不可用时回落默认
            url = (
                "postgresql+asyncpg://postgres:postgres@localhost:5432/audit_platform"
            )
    return url.replace("+asyncpg", "").replace("+psycopg2", "")


def _query(sql: str) -> list[tuple[Any, ...]]:
    """只读查询。连不上就 skip 并写明原因 —— 不静默转绿。"""
    try:
        import psycopg2
    except ImportError:  # pragma: no cover
        pytest.skip("psycopg2 未安装 ⇒ 真库判据无法取证（不视为通过）")
    try:
        conn = psycopg2.connect(_dsn(), connect_timeout=5)
    except Exception as exc:  # pragma: no cover - 无库环境
        pytest.skip(f"真库 audit_platform 连不上 ⇒ 真库判据无法取证（不视为通过）：{exc}")
    try:
        with conn, conn.cursor() as cur:
            cur.execute(sql)
            return list(cur.fetchall())
    finally:
        conn.close()


def _manifest_entries() -> list[dict[str, Any]]:
    raw = json.loads(MANIFEST_PATH.read_bytes().decode("utf-8"))
    ents = raw["entries"]
    return list(ents) if isinstance(ents, list) else list(ents.values())


def _entry(entry_id: str) -> dict[str, Any]:
    for e in _manifest_entries():
        if e.get("entry_id") == entry_id:
            return e
    raise AssertionError(f"manifest 里找不到 entry_id={entry_id!r}")


def _load_sheet(name: str):
    from openpyxl import load_workbook

    wb = load_workbook(L1_TEMPLATE, data_only=False)
    assert name in wb.sheetnames, f"{name!r} 不在册内：{wb.sheetnames}"
    return wb, wb[name]


def _is_formula(value: Any) -> bool:
    return isinstance(value, str) and value.startswith("=")


# ═══════════════════════════════════════════════════════════════════════════
# LR-P7 ~ LR-P10：受管表 `明细表L1-2` 几何基线
# ═══════════════════════════════════════════════════════════════════════════

MANAGED_GEOMETRY: dict[str, Any] = {
    "dims": "A1:AD49",
    "max_row": 49,
    "max_col": 30,
    "header_group_row": 8,
    "header_leaf_row": 9,
    "header_rows": 2,
    "first_data_row": 10,
    "last_data_row": 25,
    "footer_row": 26,
    "last_effective_col": "AB",
}
MANAGED_FORMULA_COLUMNS: tuple[str, ...] = ("K", "R", "S", "T", "U")
MANAGED_FORMULA_SHAPES: dict[str, str] = {
    "K": "=H10+I10-J10",
    "R": "=H10+L10+M10",
    "S": "=I10+N10+P10",
    "T": "=J10+O10+Q10",
    "U": "=R10+S10-T10",
}
#: 🔴 footer R26 是**全 SUM**（H..U 连续 14 列，含 K/R/S/T/U）——
#: 首版判据按探针的截断输出（`rows[:12]`）误推成 9 列，实测打红后改正。
#: 教训：计数类判据必须取完整现算结果，截断输出不可当全集。
FOOTER_SUM_COLUMNS: tuple[str, ...] = (
    "H", "I", "J", "K", "L", "M", "N", "O", "P", "Q", "R", "S", "T", "U",
)
#: footer 标签实测落在 **C 列**、文本 `合计`（无 d6 那种 3 半角空格）。
FOOTER_LABEL_CELL: str = "C26"
FOOTER_LABEL_TEXT: str = "合计"
#: D/E/F 三列 footer 占位符（长破折号），非数据。
FOOTER_PLACEHOLDER_COLUMNS: tuple[str, ...] = ("D", "E", "F")
BARE_IF_PER_SHEET: dict[str, int] = {
    "审定表L1-1": 34,
    "附注披露信息核对（上市公司）": 28,
    "附注披露信息核对（国企）": 20,
    "利息测算表L1-5": 22,
    "逾期贷款检查表L1-7": 8,
}
BARE_IF_TOTAL: int = 112


def _count_bare_if(ws) -> int:
    n = 0
    for row in ws.iter_rows():
        for c in row:
            v = c.value
            if _is_formula(v) and "IF(" in str(v).upper() and "IFERROR" not in str(v).upper():
                n += 1
    return n



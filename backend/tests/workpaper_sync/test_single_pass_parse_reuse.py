# -*- coding: utf-8 -*-
"""任务 8 · P4：materialize / extract / verify 共用一个 `workbook_read_scope()`。

spec: oo-single-pass-materialize-and-room-leave · Requirement 2.1 / 2.2

═══ 命题 ═══════════════════════════════════════════════════════════════════════

一次 CPU 段（materialize → extract → roundtrip → unmanaged verify → structure_hash）
里，**同一份产物字节**的 workbook 解析次数与「有几个消费方要读它」**解耦**：每个
解析视图各一次，不是每个消费方各一次。

与趟数判据（P1，`test_single_pass_materialize.py`）分文件：那边问「写了几趟」，
这边问「同一份字节读了几遍」。与等值判据（P2，`test_single_pass_artifact_equivalence.py`）
也分文件：那边问「提速是否以少写为代价」，这边问「复用解析是否以换结论为代价」。

═══ 「解析次数为 1」的可达下界：按**视图**算，不是按文件算 ═══════════════════════

需求 2.1 原写「对同一份产物字节的 workbook 解析次数 SHALL 为 1」，而需求 2.2 明说
`data_only=True/False` 是两个**不可互相冒充**的解析结果、作用域键里必须保留它。两条按
字面打架：只要两个视图都被用到，下界就不可能是 1。任务 7（design 附录 E.3）把这个冲突
登记为「留给任务 8 拍板」，本任务按 A.7 的同一条处理方式收紧了需求 2.1 的措辞（见
design 附录 F），判据按**收紧后**的口径写：

    每个「解析视图」= (字节身份, read_only, data_only)，每个视图各解析 **1** 次。

真库 D4 的产物字节上，实测被用到的视图恰好 3 个：

| 视图 | 谁要它 | 为什么不能与别的视图合并 |
|---|---|---|
| `read_only=True, data_only=False` | `_read_cell_view`（39 binding 读公式） | 公式文本 |
| `read_only=True, data_only=True` | `_read_cell_view`（39 binding 读公式值） | 缓存值；需求 2.2 点名 |
| `read_only=False, data_only=False` | 转置 sheet 的 8 个消费方 | `read_only` 视图没有 `ws._cells` / `row_dimensions` |

⇒ 产物字节的解析次数：改动前 **10**（2 + 8），改动后 **3**（2 + 1）。8 个消费方共用 1 次。

═══ 复用的是解析，**不是**结论 ══════════════════════════════════════════════════

需求 2.3 与 design 六 第 3 条的底线在本文件里是**可执行判据**，不是注释里的自我声明：

* `resolve_managed_sheet` 的全部校验（definedName 唯一 / workbook-scope / RANGE /
  几何漂移 / 受管区范围）逐次重跑 —— 共享的只有 `load_workbook` 的产物；
* `test_sharing_changes_no_conclusion_*` 在**同一个 world** 上把整段跑两遍（开/关复用），
  逐字段比 extract projection、逐 binding 比 verify 结论、比 structure_hash、比产物字节；
* `test_extract_and_verify_are_still_two_independent_calls` 钉住两段没被合并。

═══ 为什么不新引入第二套缓存（需求 2.2 的硬约束）═══════════════════════════════

字节身份的条目挂在**同一个** `workbook_read_scope()` 的 dict 上。任务 6（design 附录 D.2）
已经用承重反证证过为什么不能长存：把 `release_scoped_workbooks` 换成 no-op，真库链式路径
当场 `PermissionError WinError 32` 并留下 38 个删不掉的临时文件。本文件补两条结构判据：
字节条目不会被按路径释放误清、作用域退出时字节条目一个不剩。
"""
from __future__ import annotations

import contextlib
import hashlib
import inspect
import io
from pathlib import Path
from typing import Any, Iterator

import pytest
from hypothesis import HealthCheck, given, settings
from hypothesis import strategies as st

from app.services.excel_structure_fingerprint import clear_structure_fingerprint_cache
from app.services.workpaper_sync import excel_extract as EE
from app.services.workpaper_sync import phase5_d4_revenue_detail as D4
from app.services.workpaper_sync import phase5_transposed_sheet as PT
from app.services.workpaper_sync.content_mutation import ContentMutationService
from app.services.workpaper_sync.excel_extract import (
    release_scoped_workbooks,
    shared_workbook_from_bytes,
    workbook_read_scope,
)
from app.services.workpaper_sync.publish_time_structure_hash import (
    anchors_from_instrumentation_specs,
    compute_structure_hash_from_artifact,
)
from app.services.workpaper_sync.transposed_registry import resolve_transposed_specs
from tests.workpaper_sync.d4_materialize_harness import (
    D4World,
    MaterializeCallCounter,
    build_world,
)
from tests.workpaper_sync.test_single_pass_artifact_equivalence import (
    compare_zip_entries,
    projection_differences,
)

# ═══════════════════════════════════════════════════════════════════════════
# 实测常数（改这些数必须同时刷新 design 附录 F 与证据文件）
# ═══════════════════════════════════════════════════════════════════════════

#: 真库 D4 的转置 sheet 数（D4-29 客户结构 + D4-12 合同）。
TRANSPOSED_SPECS = 2

#: 产物字节上**可共享**的完整 DOM 消费方个数（每个都曾各自 `load_workbook` 一次）：
#: extract 2（每 spec 一次）+ unmanaged verify 2 + structure_hash 4（每 spec resolve + extract）。
ARTIFACT_FULL_DOM_CONSUMERS = 8

#: 产物的 `read_only` 视图数（`data_only` False/True 各一个，39 个 binding 共享）。
ARTIFACT_READ_ONLY_VIEWS = 2

#: 产物字节上**不可共享**的完整 DOM 解析，逐个列名 + 理由。诚实地留在判据里而不是从
#: 「解析次数」里抹掉：requirements 2.1 问的是「同一份字节读了几遍」，把读不掉的那几遍
#: 藏起来就是把判据写成自己想要的答案。
#:
#: 两者都用 openpyxl 的**批量成格**读法（`ws.iter_rows()` 会把包围盒里每个坐标都惰性
#: 新建成 Cell），共享给它们等于污染 `ws._cells` ⇒ `collect_workbook_structure` 的
#: 「这个转置字段格在文件里存不存在」当场变答案。所以它们**必须**各自私有解析。
ARTIFACT_EXCLUSIVE_FULL_DOM_SITES = (
    # 整簿指纹：`iter_rows()` 遍历每张 sheet 的每个格子取值/公式/样式。
    "_structure_fingerprint_uncached",
    # 隐藏 metadata sheet 的 key→value：`iter_rows(min_col=1, max_col=2)`。
    "_read_gt_sync_pairs",
)

#: 整段 `load_workbook` 的**完整**调用点清册（顺序 = 报告顺序）。
SEGMENT_LOAD_SITES = (
    "resolve_managed_sheet",
    "shared_workbook_from_bytes",
    "_acquire_read_only_workbook",
    *ARTIFACT_EXCLUSIVE_FULL_DOM_SITES,
)

#: 逐调用点实测（任务 7 的 design 附录 E.3 给了「前」这一列）。
SEGMENT_LOAD_CENSUS_WITHOUT_REUSE = {
    # materialize 私有 4（每张转置 sheet 读一次 + 写一次，中间字节各不相同）
    # + 产物 8（extract 2 / verify 2 / structure_hash 4，各自重解一遍）
    "resolve_managed_sheet": 12,
    "shared_workbook_from_bytes": 0,
    # substrate 两视图（39 binding 共享）+ 产物两视图（39 binding 共享）
    "_acquire_read_only_workbook": 4,
    # substrate 侧 1 + 产物侧 1（内容寻址记忆化，段内两份字节各一次）
    "_structure_fingerprint_uncached": 2,
    "_read_gt_sync_pairs": 1,
    "<其它>": 0,
}
SEGMENT_LOAD_CENSUS_WITH_REUSE = {
    **SEGMENT_LOAD_CENSUS_WITHOUT_REUSE,
    "resolve_managed_sheet": 4,
    "shared_workbook_from_bytes": 1,
}

#: 整段 `load_workbook` 总数（任务 7 实测 19 —— design 附录 E.3 的逐调用点清单）。
SEGMENT_LOADS_WITHOUT_REUSE = 19
SEGMENT_LOADS_WITH_REUSE = 12

#: 生产 CPU 段的五步标记（与 `scripts/analyze/measure_d4_materialize_baseline.py` 同一份口径）。
SEGMENT_STEP_MARKERS = (
    "adapter.materialize(",
    "adapter.extract(",
    "self._assert_roundtrip_equivalent(",
    "adapter.verify_unmanaged_regions(",
    "self._projection_structure_hash(",
)


# ═══════════════════════════════════════════════════════════════════════════
# 观测器（一律只观测，不改行为）
# ═══════════════════════════════════════════════════════════════════════════


def _digest(data: bytes) -> str:
    return hashlib.sha256(data).hexdigest()


def cells_digest(worksheet: Any) -> str:
    """一张 sheet 的 `_cells` **键集合**摘要 —— 「有没有人就地改过它」的 oracle。

    openpyxl 的 `ws.cell(r, c)` 在格子缺失时会新建一个空 Cell 并塞进 `_cells`，那是
    就地改动；而 `collect_workbook_structure` 正是按 `(row, col) not in ws._cells` 判
    转置字段格在文件里存不存在 ⇒ 键集合变了就等于结论变了。比键集合而不是比值：
    本判据要抓的正是「多出了本不存在的格子」。
    """
    keys = getattr(worksheet, "_cells", None)
    if keys is None:  # pragma: no cover - read_only 视图不走本判据
        return "<no-_cells>"
    return hashlib.sha256(repr(sorted(keys)).encode("utf-8")).hexdigest()


class ResolveCall:
    """一次 `resolve_managed_sheet` 调用的观测：读了谁的字节、共享没共享、**是否真解析了**。

    `fresh` 不靠「数调用点行号出现几次」推：直接看这次调用之后作用域缓存里有没有多出
    一个键。命中共享 ⇒ 键数不变 ⇒ `fresh=False`；私有解析或首次进缓存 ⇒ `fresh=True`。
    这样「8 个消费方共用 1 次解析」是**从缓存行为**读出来的，不是从行号统计里猜的。

    `cells` 是**调用返回那一刻**受管 sheet 的 `_cells` 键集合摘要。连续两次共享调用之间
    夹着一个消费方，两次摘要不同 ⇒ 那个消费方改了共享对象。
    """

    __slots__ = ("sheet_key", "digest", "shared", "fresh", "sheet_title", "cells")

    def __init__(
        self,
        *,
        sheet_key: str,
        digest: str,
        shared: bool,
        fresh: bool,
        sheet_title: str,
        cells: str,
    ) -> None:
        self.sheet_key = sheet_key
        self.digest = digest
        self.shared = shared
        self.fresh = fresh
        self.sheet_title = sheet_title
        self.cells = cells

    def __repr__(self) -> str:  # pragma: no cover - 只在断言失败消息里出现
        flag = "shared" if self.shared else "private"
        return (
            f"{self.sheet_key}@{self.digest[:12]} {flag} "
            f"{'parse' if self.fresh else 'hit'} cells={self.cells[:8]}"
        )


@contextlib.contextmanager
def resolve_witness() -> Iterator[list[ResolveCall]]:
    """记录每次 `resolve_managed_sheet(bytes, spec, share_parse)`。

    只打 `phase5_transposed_sheet.resolve_managed_sheet` 一处就够：`adapters.excel` 与
    `published_identity_observer` 都在**函数体内** import（每次调用重取模块属性），而
    `extract_transposed_workbook` 在本模块内按全局名解析它 ⇒ 这一处覆盖全部调用点。
    观测器自身的非空转由 `test_cpu_segment_parses_the_artifact_bytes_once_per_view` 的
    「必须记到 8 次调用」承担（记不到就说明观测器失效，而不是「没人读产物」）。
    """
    seen: list[ResolveCall] = []
    original = PT.resolve_managed_sheet

    def watched(workbook_bytes: bytes, *, spec: Any, share_parse: bool = False) -> Any:
        before = len(scoped_keys())
        in_scope = EE._workbook_scope.get() is not None
        result = original(workbook_bytes, spec=spec, share_parse=share_parse)
        grew = len(scoped_keys()) > before
        _wb, worksheet = result
        seen.append(
            ResolveCall(
                sheet_key=str(spec.sheet_key),
                digest=_digest(workbook_bytes),
                shared=bool(share_parse),
                # 不共享 / 不在作用域内 ⇒ 必然当场解析一份私有的。
                fresh=(not share_parse) or (not in_scope) or grew,
                sheet_title=str(worksheet.title),
                cells=cells_digest(worksheet),
            )
        )
        return result

    PT.resolve_managed_sheet = watched  # type: ignore[assignment]
    try:
        yield seen
    finally:
        PT.resolve_managed_sheet = original  # type: ignore[assignment]


#: 被记忆化的三处**纯 XML/zip 解析**：`调用次数` vs `真解析次数`。
XML_PARSE_PROBES = (
    # (计数名, 包装函数名, 本体函数名)
    ("workbook_xml", "_sheet_part_map", "parse_workbook_xml"),
    ("tables", "_tables_of", "parse_tables"),
    ("shared_strings", "_shared_strings_prefix_digest", "_shared_strings_prefix_digest_uncached"),
)


@contextlib.contextmanager
def xml_parse_witness() -> Iterator[dict[str, int]]:
    """对三处记忆化解析各记「被调了几次」与「真解析了几次」。

    全部打在 `excel_extract` 的**模块属性**上：生产里 `resolve_managed_region` /
    `unmanaged_region_digest` 都按模块全局名解析它们 ⇒ 打模块就打得到（与
    `MaterializeCallCounter` 对 `load_workbook` 的做法同一条纪律）。
    """
    tally: dict[str, int] = {}
    undo: list[Any] = []
    for label, wrapper, inner in XML_PARSE_PROBES:
        tally[f"{label}_calls"] = 0
        # 真解析分两类：**有文件身份**的（可记忆化）与 `BytesIO` 上打开的 zip（拿不到
        # 身份 ⇒ 按 `scoped_parse_memo` 的约定退化成不记忆化）。混成一个数会让「记忆化
        # 生效了没有」读不出来。
        tally[f"{label}_parses_identified"] = 0
        tally[f"{label}_parses_anonymous"] = 0

        original_wrapper = getattr(EE, wrapper)

        def watched_wrapper(
            *args: Any, _o: Any = original_wrapper, _k: str = f"{label}_calls", **kwargs: Any
        ) -> Any:
            tally[_k] += 1
            return _o(*args, **kwargs)

        undo.append(lambda n=wrapper, o=original_wrapper: setattr(EE, n, o))
        setattr(EE, wrapper, watched_wrapper)

        original_inner = getattr(EE, inner)

        def watched_inner(
            zf: Any, *args: Any, _o: Any = original_inner, _label: str = label, **kwargs: Any
        ) -> Any:
            identified = EE._zip_identity(zf) is not None
            suffix = "identified" if identified else "anonymous"
            tally[f"{_label}_parses_{suffix}"] += 1
            return _o(zf, *args, **kwargs)

        undo.append(lambda n=inner, o=original_inner: setattr(EE, n, o))
        setattr(EE, inner, watched_inner)
    try:
        yield tally
    finally:
        for restore in reversed(undo):
            restore()


@contextlib.contextmanager
def reuse_disabled() -> Iterator[None]:
    """承重反证：把 `share_parse` 一律掐成 `False`（= 改动前的形态）。

    刻意不改判据、不改调用点、不改 adapter —— 只在生产内核入口把那个开关摘掉，
    于是「8 个消费方各解析一次」这个改动前的事实会**真的**重现。留着复用跑一遍全绿
    只说明「没坏」，说明不了复用有用。
    """
    original = PT.resolve_managed_sheet

    def without_sharing(workbook_bytes: bytes, *, spec: Any, share_parse: bool = False) -> Any:
        return original(workbook_bytes, spec=spec, share_parse=False)

    PT.resolve_managed_sheet = without_sharing  # type: ignore[assignment]
    try:
        yield
    finally:
        PT.resolve_managed_sheet = original  # type: ignore[assignment]


def scoped_keys() -> tuple[tuple[Any, ...], ...]:
    """当前作用域缓存的全部键（确定顺序）。直接读生产 `ContextVar`，不另算一套。"""
    cache = EE._workbook_scope.get() or {}
    return tuple(sorted(cache, key=repr))


def byte_scoped_keys() -> tuple[tuple[Any, ...], ...]:
    """其中**字节身份**的条目。"""
    return tuple(
        key
        for key in scoped_keys()
        if isinstance(key[0], tuple) and key[0][0] == EE._BYTES_SCOPE_TAG
    )


def path_scoped_views(path: Path) -> tuple[bool, ...]:
    """作用域缓存里针对 `path` 的 `data_only` 视图集合 —— 与 `release_scoped_workbooks`
    的匹配口径同一条（`k[0][0] == str(path.resolve())`）。"""
    cache = EE._workbook_scope.get() or {}
    target = str(path.resolve())
    return tuple(
        sorted(
            bool(key[1])
            for key in cache
            if isinstance(key[0], tuple) and key[0][0] == target
        )
    )


# ═══════════════════════════════════════════════════════════════════════════
# 生产 CPU 段的复刻（形态自检在 `test_the_replica_matches_the_production_segment`）
# ═══════════════════════════════════════════════════════════════════════════


class SegmentRun:
    """跑一次 CPU 段的全部观测量 + 全部**结论**。"""

    def __init__(
        self,
        *,
        counter: MaterializeCallCounter,
        resolves: list[ResolveCall],
        artifact: Path,
        artifact_digest: str,
        extracted: Any,
        verify_verdicts: list[bool],
        verify_first_difference: Any,
        structure_hash: str,
        byte_keys_at_end: tuple[tuple[Any, ...], ...],
        artifact_path_views: tuple[bool, ...],
        cells_after_segment: dict[str, str],
        parse_memo: dict[tuple[Any, ...], str],
        xml_parses: dict[str, int],
    ) -> None:
        self.cells_after_segment = cells_after_segment
        self.parse_memo = parse_memo
        self.xml_parses = xml_parses
        self.counter = counter
        self.resolves = resolves
        self.artifact = artifact
        self.artifact_digest = artifact_digest
        self.extracted = extracted
        self.verify_verdicts = verify_verdicts
        self.verify_first_difference = verify_first_difference
        self.structure_hash = structure_hash
        self.byte_keys_at_end = byte_keys_at_end
        self.artifact_path_views = artifact_path_views

    @property
    def artifact_resolve_calls(self) -> tuple[ResolveCall, ...]:
        """`resolve_managed_sheet` 中读**最终产物字节**的那些调用。"""
        return tuple(call for call in self.resolves if call.digest == self.artifact_digest)

    @property
    def artifact_full_dom_parses(self) -> int:
        """产物字节上**真正**发生的完整 DOM 解析次数（命中共享的不算）。"""
        return sum(1 for call in self.artifact_resolve_calls if call.fresh)

    def loads_at(self, fragment: str) -> int:
        """调用点含 `fragment` 的 `load_workbook` 次数（分调用点的分母）。"""
        return sum(
            count
            for site, count in self.counter.load_sites.items()
            if fragment in site
        )

    def load_census(self) -> dict[str, int]:
        """整段 `load_workbook` 按调用点归并 —— **完整**清册，一次都不许落在册外。

        「总数从 19 降到 12」这句话单独看可以糊过去（谁知道降的是哪几次）。清册逐项对账，
        每一项都得有名字与理由，新出现的调用点会让 `<其它>` 非零而打红。
        """
        census = {name: self.loads_at(name) for name in SEGMENT_LOAD_SITES}
        census["<其它>"] = self.counter.load_count - sum(census.values())
        return census


@contextlib.contextmanager
def _collecting_unmanaged_verdicts(sink: list[bool]) -> Iterator[None]:
    """把 `UnmanagedRegionReport.assert_equivalent` 换成收集器（只在 verify 段内）。

    与任务 7（design 附录 E.2）同一个理由：生产那个循环**逐 binding 比完就 assert**，
    harness world 的第 1 个 binding 就报 unmanaged drift（模板 → 第一代产物造成，
    steady-state 探针实测 39/39 等价）⇒ 不换成收集器的话循环只跑 39 分之一，
    「开/关复用结论相同」这条判据就只比了一个 binding。结论一条不藏，逐个进断言。
    """
    from app.services.workpaper_sync.adapters.base import UnmanagedRegionReport

    original = UnmanagedRegionReport.assert_equivalent

    def collecting(self: Any) -> None:
        sink.append(bool(self.equivalent))

    UnmanagedRegionReport.assert_equivalent = collecting  # type: ignore[method-assign]
    try:
        yield
    finally:
        UnmanagedRegionReport.assert_equivalent = original  # type: ignore[method-assign]


def run_cpu_segment(world: D4World, *, output_name: str, reuse: bool) -> SegmentRun:
    """在**一个** `workbook_read_scope()` 里跑生产 CPU 段五步（与生产同形）。"""
    # 结构指纹按字节内容记忆化 ⇒ 不清的话第二次运行白捡第一次的钱，
    # 「开/关复用」的解析计数就不可比了。
    clear_structure_fingerprint_cache()
    anchors = anchors_from_instrumentation_specs(D4.instrumentation_specs())
    output = world.staged(output_name)
    counter = MaterializeCallCounter()
    verdicts: list[bool] = []
    cells_after: dict[str, str] = {}
    gate = contextlib.nullcontext() if reuse else reuse_disabled()
    # 🔴 顺序有意义：观测器**先**进、`reuse_disabled()` **后**进 ⇒ 调用先被掐掉
    # `share_parse` 再被观测，观测到的是**实际发生**的那次调用（而不是调用方的意图）。
    # 反过来的话，反证侧会被记成「带着 share_parse 却没命中缓存」—— 那读起来像复用失效，
    # 实际上是观测器站错了位置。
    with resolve_witness() as resolves, gate, xml_parse_witness() as xml_parses, workbook_read_scope():
        with counter.installed():
            materialized = world.adapter.materialize(
                substrate=world.base,
                projection=world.projection,
                output=output,
                contract=world.contract,
            )
            extracted = world.adapter.extract(artifact=output, contract=world.contract)
            ContentMutationService._assert_roundtrip_equivalent(  # noqa: SLF001
                None,  # type: ignore[arg-type]
                intended=world.projection,
                extracted=extracted,
                contract=world.contract,
            )
            with _collecting_unmanaged_verdicts(verdicts):
                report = world.adapter.verify_unmanaged_regions(
                    before=world.base,
                    after=output,
                    contract=world.contract,
                    row_shift=materialized.row_shift,
                    total_formula_rows=materialized.total_formula_rows,
                    propagation=materialized.workbook_row_change,
                    per_table_shift=materialized.per_table_shift,
                )
            structure_hash = compute_structure_hash_from_artifact(
                data=output.read_bytes(), contract=world.contract, anchors=anchors
            )
        # 🔴 快照必须在作用域**里**取：退出时条目会被关掉并清空。
        byte_keys = byte_scoped_keys()
        artifact_views = path_scoped_views(output)
        cache = EE._workbook_scope.get() or {}
        parse_memo = {
            key[0]: repr(value)
            for key, value in cache.items()
            if len(key) == 1
            and isinstance(key[0], tuple)
            and key[0][0] == EE._PARSE_MEMO_TAG
        }
        if reuse:
            # 最后一个消费方之后的状态 —— 计数桩已退出，这次命中缓存也不会污染计数
            # （反证侧不取：那边压根没有共享对象，`shared_workbook_from_bytes` 会白解析一份）。
            shared = shared_workbook_from_bytes(output.read_bytes())
            for title in sorted(
                {
                    call.sheet_title
                    for call in resolves
                    if call.digest == _digest(output.read_bytes())
                }
            ):
                cells_after[title] = cells_digest(shared[title])
    return SegmentRun(
        counter=counter,
        resolves=resolves,
        artifact=output,
        artifact_digest=_digest(output.read_bytes()),
        extracted=extracted,
        verify_verdicts=verdicts,
        verify_first_difference=report.first_difference,
        structure_hash=structure_hash,
        byte_keys_at_end=byte_keys,
        artifact_path_views=artifact_views,
        cells_after_segment=cells_after,
        parse_memo=parse_memo,
        xml_parses=dict(xml_parses),
    )


@pytest.fixture(scope="module")
def world(tmp_path_factory: pytest.TempPathFactory) -> D4World:
    """真 D4 模板 + 真契约 + 真 39 binding；module 作用域（铺一次 world 十几秒）。"""
    return build_world(tmp_path_factory.mktemp("d4-parse-reuse"))


@pytest.fixture(scope="module")
def with_reuse(world: D4World) -> SegmentRun:
    return run_cpu_segment(world, output_name="reuse-on.xlsx", reuse=True)


@pytest.fixture(scope="module")
def without_reuse(world: D4World) -> SegmentRun:
    return run_cpu_segment(world, output_name="reuse-off.xlsx", reuse=False)


# ═══════════════════════════════════════════════════════════════════════════
# 一、复刻保真 + 前提（分母在这里；没有它们下面的计数都可能是空转）
# ═══════════════════════════════════════════════════════════════════════════


def test_the_replica_matches_the_production_segment() -> None:
    """**Validates: Requirements 2.1**

    本文件量的是「生产 CPU 段」。生产那段的调用清单一变，复刻就过期了 —— 那时下面所有
    解析计数都是在量一段**不存在**的代码。所以先对账：五步逐条在生产源码里找得到，且
    生产确实用**一个** `workbook_read_scope()` 包住整段（需求 2.2 说的复用载体）。
    """
    scoped = inspect.getsource(ContentMutationService._stage_cpu_segment_scoped)
    outer = inspect.getsource(ContentMutationService._stage_cpu_segment)
    missing = [marker for marker in SEGMENT_STEP_MARKERS if marker not in scoped]
    assert not missing, (
        f"生产 CPU 段里找不到这些步骤：{missing} ⇒ 复刻已过期，本文件的计数不再代表生产"
    )
    assert "with workbook_read_scope():" in outer, (
        "生产 CPU 段不再用 `workbook_read_scope()` 包住整段 ⇒ 需求 2.2 的复用载体没了，"
        "本文件测到的复用只发生在测试自己开的作用域里"
    )
    assert "self._session" not in scoped, (
        "CPU 段开始碰 DB 会话了（Property 4 / Requirement 2.2 禁止）"
    )


def test_the_world_really_has_the_transposed_sheets_this_judgement_is_about(
    world: D4World,
) -> None:
    """**Validates: Requirements 2.1**

    产物字节上那 8 个完整 DOM 消费方全部来自转置 sheet。world 上要是一张转置 sheet 都没有，
    「8 次合成 1 次」这条判据的分母就是 0，它会永远绿。
    """
    specs = resolve_transposed_specs(world.contract)
    assert len(specs) == TRANSPOSED_SPECS, (
        f"转置 spec 数 {len(specs)} ≠ 实测 {TRANSPOSED_SPECS}"
        f"（{[s.sheet_key for s in specs]}）—— 契约的转置声明变了，"
        "本文件的常数与 design 附录 F 必须一起刷新"
    )


# ═══════════════════════════════════════════════════════════════════════════
# 二、P4：同一份字节，每个视图各解析一次（需求 2.1）
# ═══════════════════════════════════════════════════════════════════════════


def test_cpu_segment_parses_the_artifact_bytes_once_per_view(
    with_reuse: SegmentRun,
) -> None:
    """**Validates: Requirements 2.1**　**Property: P4**

    本任务的**主判据**。三件事一起断，缺一件都能被糊过去：

    1. 产物字节上确实有 8 个完整 DOM 消费方在读（分母 —— 观测器没失效、路径没改道）；
    2. 这 8 次里**只有 1 次**真解析（其余命中同一个作用域缓存条目）；
    3. 算上 `_read_cell_view` 的两个 `read_only` 视图，产物字节的解析总数 = 3
       —— 即「每个视图各 1 次」，与消费方个数解耦。
    """
    calls = with_reuse.artifact_resolve_calls
    assert len(calls) == ARTIFACT_FULL_DOM_CONSUMERS, (
        f"读产物字节的 resolve_managed_sheet 调用 {len(calls)} 次 ≠ 实测 "
        f"{ARTIFACT_FULL_DOM_CONSUMERS}（extract 2 + verify 2 + structure_hash 4）；"
        f"实际清册：{list(calls)}"
    )
    assert all(call.shared for call in calls), (
        "有消费方没开 share_parse ⇒ 它自己又解析了一遍产物："
        f"{[call for call in calls if not call.shared]}"
    )
    assert with_reuse.artifact_full_dom_parses == 1, (
        f"产物字节的完整 DOM 解析 {with_reuse.artifact_full_dom_parses} 次（应为 1）；"
        f"清册：{list(calls)}"
    )
    assert with_reuse.artifact_path_views == (False, True), (
        f"产物的 read_only 视图 {with_reuse.artifact_path_views} ≠ (False, True) —— "
        "要么 39 个 binding 没共用那两个视图，要么两个视图被串成了一条（需求 2.2 禁止）"
    )
    assert len(with_reuse.artifact_path_views) == ARTIFACT_READ_ONLY_VIEWS


def test_without_sharing_every_consumer_parses_the_artifact_again(
    without_reuse: SegmentRun,
) -> None:
    """**Validates: Requirements 2.1**　**Property: P4 的反证**

    承重反证（design 五 P4 的「去掉 scope 复用 ⇒ 解析计数 >1 ⇒ 红」）：把 `share_parse`
    掐成 `False`，8 个消费方必须**真的**各解析一遍 —— 这就是改动前的形态。这一条红了
    说明上一条不是空转；这一条绿而上一条也绿，说明复用真的把 8 变成了 1。
    """
    calls = without_reuse.artifact_resolve_calls
    assert len(calls) == ARTIFACT_FULL_DOM_CONSUMERS, (
        f"反证侧读产物字节 {len(calls)} 次（应 {ARTIFACT_FULL_DOM_CONSUMERS}）"
    )
    assert without_reuse.artifact_full_dom_parses == ARTIFACT_FULL_DOM_CONSUMERS, (
        f"掐掉复用后产物字节只解析了 {without_reuse.artifact_full_dom_parses} 次"
        f"（应 {ARTIFACT_FULL_DOM_CONSUMERS}）⇒ `reuse_disabled()` 没生效，"
        "主判据的对照分母不成立"
    )
    assert not any(call.shared for call in calls), "反证侧仍有调用带着 share_parse"


def test_segment_load_workbook_total_drops_by_the_shared_parses(
    with_reuse: SegmentRun, without_reuse: SegmentRun
) -> None:
    """**Validates: Requirements 2.1**

    整段总数（任务 7 的 design 附录 E.3 实测 19）必须**按差额**下降，且下降全部来自
    产物字节那 7 次：materialize 自己那 4 次（读的是每张转置 sheet 写入**之前**的中间
    字节，各不相同、写完再没人读）刻意不共享，总数里必须还在。
    """
    assert without_reuse.counter.load_count == SEGMENT_LOADS_WITHOUT_REUSE, (
        f"复用前整段 load_workbook {without_reuse.counter.load_count} ≠ 基线 "
        f"{SEGMENT_LOADS_WITHOUT_REUSE}；调用点：{dict(without_reuse.counter.load_sites)}"
    )
    assert with_reuse.counter.load_count == SEGMENT_LOADS_WITH_REUSE, (
        f"复用后整段 load_workbook {with_reuse.counter.load_count} ≠ 实测 "
        f"{SEGMENT_LOADS_WITH_REUSE}；调用点：{dict(with_reuse.counter.load_sites)}"
    )
    # 清册对账：省下的到底是哪几次、剩下的每一次是谁，全部列名（含「不可共享」那两个）。
    assert without_reuse.load_census() == SEGMENT_LOAD_CENSUS_WITHOUT_REUSE, (
        f"复用前的调用点清册变了：{without_reuse.load_census()} != "
        f"{SEGMENT_LOAD_CENSUS_WITHOUT_REUSE}"
    )
    assert with_reuse.load_census() == SEGMENT_LOAD_CENSUS_WITH_REUSE, (
        f"复用后的调用点清册变了：{with_reuse.load_census()} != "
        f"{SEGMENT_LOAD_CENSUS_WITH_REUSE}"
    )
    saved = without_reuse.counter.load_count - with_reuse.counter.load_count
    assert saved == ARTIFACT_FULL_DOM_CONSUMERS - 1, (
        f"省下 {saved} 次，应为 {ARTIFACT_FULL_DOM_CONSUMERS - 1} 次"
    )
    # materialize 的私有解析（转置中间字节）两侧都得在 —— 它不该被「顺手」共享掉。
    materialize_private = [
        call for call in with_reuse.resolves if call.digest != with_reuse.artifact_digest
    ]
    assert len(materialize_private) == TRANSPOSED_SPECS * 2, (
        f"materialize 侧转置解析 {len(materialize_private)} 次 ≠ "
        f"{TRANSPOSED_SPECS * 2}（每 spec 读一次 + 写一次）：{materialize_private}"
    )
    assert all(not call.shared and call.fresh for call in materialize_private), (
        "materialize 侧的转置解析被共享了 —— 其中写入那次会就地改 workbook + `wb.save()`，"
        f"共享等于把改动串给 verify / structure_hash：{materialize_private}"
    )


def _tiny_workbook_bytes(*, formula: str = "=1+1", cached: float = 2.0) -> bytes:
    """一份最小 xlsx：A1 是公式、缓存值已写好 ⇒ 两个 `data_only` 视图读出不同的东西。

    刻意不用真库 D4：本组判据问的是**作用域机制**本身（键里含不含视图、会不会互相冒充），
    用几 KB 的字节跑得快、也不受契约变化影响。视图差异的**存在**由下面的断言实测，
    不是假设。
    """
    import openpyxl

    workbook = openpyxl.Workbook()
    sheet = workbook.active
    sheet["A1"] = formula
    buffer = io.BytesIO()
    workbook.save(buffer)
    raw = buffer.getvalue()
    # openpyxl 不写缓存值 ⇒ 手工把 `<f>` 旁边的 `<v>` 补上，否则 data_only=True 读出 None，
    # 「两视图不同」就只是「一个有值一个没值」的弱证据。
    import re
    import zipfile

    out = io.BytesIO()
    with zipfile.ZipFile(io.BytesIO(raw)) as src, zipfile.ZipFile(
        out, "w", zipfile.ZIP_DEFLATED
    ) as dst:
        for info in src.infolist():
            data = src.read(info.filename)
            if info.filename == "xl/worksheets/sheet1.xml":
                data = re.sub(
                    rb"(<f>[^<]*</f>)", rb"\1<v>" + str(cached).encode() + rb"</v>", data
                )
            dst.writestr(info, data)
    return out.getvalue()


# ═══════════════════════════════════════════════════════════════════════════
# 三、需求 2.2：复用走既有作用域，且视图之间不得互相冒充
# ═══════════════════════════════════════════════════════════════════════════


def test_data_only_views_of_the_same_bytes_are_separate_entries() -> None:
    """**Validates: Requirements 2.2**

    design 三点名的那条：`data_only=True/False` 是**不同**的解析结果，不可互相冒充。
    两面都断：
    * 结构 —— 两次 acquire 落在两个不同的缓存键上、拿到两个不同的对象；
    * 后果 —— 同一个格子在两个视图里读出**不同**的东西（公式文本 vs 缓存值）。
      只断结构的话，「键分开了但值其实一样」这种退化（例如有人把 data_only 传丢了）
      照样绿。
    """
    data = _tiny_workbook_bytes()
    with workbook_read_scope():
        formulas = shared_workbook_from_bytes(data, data_only=False)
        values = shared_workbook_from_bytes(data, data_only=True)
        assert formulas is not values, "两个 data_only 视图拿到了同一个对象 ⇒ 互相冒充"
        assert len(byte_scoped_keys()) == 2, (
            f"两个视图只落了 {len(byte_scoped_keys())} 个缓存键：{byte_scoped_keys()}"
        )
        assert formulas["Sheet"]["A1"].value == "=1+1", (
            f"data_only=False 应读公式文本，实得 {formulas['Sheet']['A1'].value!r}"
        )
        assert values["Sheet"]["A1"].value == 2.0, (
            f"data_only=True 应读缓存值，实得 {values['Sheet']['A1'].value!r}"
        )
        # 同一视图再要一次 ⇒ 必须是同一个对象（这才是「只解析一次」）。
        assert shared_workbook_from_bytes(data, data_only=False) is formulas
        assert shared_workbook_from_bytes(data, data_only=True) is values
        assert len(byte_scoped_keys()) == 2


def test_read_only_view_never_impersonates_the_full_dom_view() -> None:
    """**Validates: Requirements 2.2**

    `read_only` 也进键，理由与 `data_only` 同一条：`read_only=True` 的 worksheet **没有**
    `_cells` / `row_dimensions` 这些完整 DOM 才有的面，而生产恰恰靠它们判「这个格子在
    文件里存不存在」（`collect_workbook_structure`）。两种解析结果互相冒充会让那条判据
    读到一个语义不同的对象。
    """
    data = _tiny_workbook_bytes()
    with workbook_read_scope():
        full = shared_workbook_from_bytes(data, read_only=False)
        lazy = shared_workbook_from_bytes(data, read_only=True)
        assert full is not lazy, "完整 DOM 与 read_only 视图拿到了同一个对象"
        assert len(byte_scoped_keys()) == 2
        assert hasattr(full["Sheet"], "_cells"), "完整 DOM 视图应有 `_cells`"
        assert not hasattr(lazy["Sheet"], "_cells"), (
            "read_only 视图居然有 `_cells` —— openpyxl 语义变了，本条的理由需重写"
        )


def test_byte_entries_live_in_the_one_existing_scope_and_die_with_it() -> None:
    """**Validates: Requirements 2.2**

    「不得新引入第二套 workbook 缓存」是需求 2.2 的硬约束（模块级长存缓存会让 Windows
    删不掉临时文件 —— 任务 6 的承重反证：release 换 no-op ⇒ WinError 32 + 38 个残留）。
    三条一起断：
    * 字节条目挂在**生产那个** `ContextVar` 上（判据直接读 `EE._workbook_scope`）；
    * 作用域之外 acquire **不进任何缓存**（退化成即用即弃，与优化前逐字节相同）；
    * 作用域退出后字节条目一个不剩。
    """
    data = _tiny_workbook_bytes()
    assert EE._workbook_scope.get() is None, "前提不成立：本条要求当前没有作用域"
    outside = shared_workbook_from_bytes(data)
    assert outside is not None
    assert EE._workbook_scope.get() is None, "作用域外的 acquire 居然建了作用域"

    with workbook_read_scope():
        first = shared_workbook_from_bytes(data)
        assert len(byte_scoped_keys()) == 1
        assert shared_workbook_from_bytes(data) is first
    assert EE._workbook_scope.get() is None
    # 退出后再进一个新作用域：缓存必须是空的（没有任何东西活过作用域）。
    with workbook_read_scope():
        assert byte_scoped_keys() == (), (
            f"新作用域里就有字节条目：{byte_scoped_keys()} ⇒ 有人把它挂到了模块级"
        )
        again = shared_workbook_from_bytes(data)
        assert again is not first, "跨作用域复用了同一个 workbook 对象 ⇒ 存在长存缓存"


def test_release_by_path_and_byte_entries_do_not_disturb_each_other(
    tmp_path: Path,
) -> None:
    """**Validates: Requirements 2.2, 1.4**

    `release_scoped_workbooks(path)` 按 `k[0][0] == str(path.resolve())` 匹配。字节条目的
    键前缀刻意不是路径形态（`bytes:sha256`）⇒ 两类条目不可能互相误清：
    * 按路径释放**不会**顺手砍掉字节条目（否则产物的共享解析会莫名其妙失效）；
    * 字节条目**不会**挡住按路径释放（否则 Windows 上临时文件又删不掉 —— 需求 1.4）。
    """
    assert not EE._BYTES_SCOPE_TAG.startswith(("/", "\\")) and ":" in EE._BYTES_SCOPE_TAG
    book = tmp_path / "release-vs-bytes.xlsx"
    book.write_bytes(_tiny_workbook_bytes())
    data = book.read_bytes()
    with workbook_read_scope():
        with EE._acquire_read_only_workbook(book, data_only=False):
            pass
        with EE._acquire_read_only_workbook(book, data_only=True):
            pass
        shared_workbook_from_bytes(data)
        assert path_scoped_views(book) == (False, True)
        assert len(byte_scoped_keys()) == 1

        release_scoped_workbooks(book)
        assert path_scoped_views(book) == (), "按路径释放没清干净"
        assert len(byte_scoped_keys()) == 1, (
            "按路径释放把字节条目也清了 ⇒ 两类键撞上了"
        )
        # 路径条目已释放 ⇒ 文件可删（Windows 上句柄泄漏时这一步抛 WinError 32）。
        book.unlink()
    assert not book.exists()


# ═══════════════════════════════════════════════════════════════════════════
# 四、复用解析 ≠ 复用结论（需求 2.3 的前提，不越界做任务 9 的变异反证）
# ═══════════════════════════════════════════════════════════════════════════


def test_sharing_changes_no_conclusion_of_extract_verify_or_structure_hash(
    with_reuse: SegmentRun, without_reuse: SegmentRun
) -> None:
    """**Validates: Requirements 2.1, 2.3**

    复用的只能是**解析结果**，不能是**结论**。同一个 world / 同一份 projection 把整段跑
    两遍（开复用 / 关复用），四个面必须逐项相同：

    * extract projection —— 用任务 4 的 `repr` 口径全字段比对（`0` / `0.0` / `Decimal('0')`
      在 `==` 下相等而落盘字节不同，用 `==` 比会让这条退化成空话）；
    * verify 的**逐 binding** 结论（39 条，不是「最后一条」）；
    * structure_hash —— 它读 `ws._cells` 的存在性，是对共享对象最敏感的那个消费方；
    * 产物字节 —— 逐 zip entry 比（`date_time` 显式分流，任务 4 附录 B.3 同一口径）。
    """
    differences = projection_differences(
        without_reuse.extracted,
        with_reuse.extracted,
        left_label="reuse_off",
        right_label="reuse_on",
    )
    assert differences == (), (
        "开/关解析复用得到的 extract projection 不同 ⇒ 复用改了结论：\n"
        + "\n".join(differences)
    )
    assert len(with_reuse.extracted.values) > 0, (
        "extract 出来 0 个字段 ⇒ 上一条的分母是空的"
    )
    assert with_reuse.verify_verdicts == without_reuse.verify_verdicts, (
        "逐 binding 的 unmanaged 结论不同 ⇒ 复用改了 verify 的答案：\n"
        f"开复用 {with_reuse.verify_verdicts}\n关复用 {without_reuse.verify_verdicts}"
    )
    assert len(with_reuse.verify_verdicts) >= TRANSPOSED_SPECS, (
        f"verify 只产出 {len(with_reuse.verify_verdicts)} 条结论 ⇒ 循环没跑完，"
        "本条比的不是完整比对面"
    )
    assert with_reuse.verify_first_difference == without_reuse.verify_first_difference, (
        f"首个差异面不同：{with_reuse.verify_first_difference} vs "
        f"{without_reuse.verify_first_difference}"
    )
    assert with_reuse.structure_hash == without_reuse.structure_hash, (
        f"structure_hash 不同：{with_reuse.structure_hash} vs "
        f"{without_reuse.structure_hash} ⇒ 共享对象被就地改过（`ws._cells` 被惰性新建污染"
        "是最可能的原因），发布时刻的结构身份因此漂了"
    )
    comparison = compare_zip_entries(without_reuse.artifact, with_reuse.artifact)
    assert comparison.namelist_equal and not comparison.name_differences, (
        f"两侧产物的 zip 条目集合/顺序不同：{comparison.name_differences}"
    )
    assert comparison.content_differences == (), (
        f"两侧产物内容不同的条目：{comparison.content_differences}"
    )
    assert comparison.entry_count > 0


def test_no_read_only_consumer_mutates_the_shared_workbook(
    with_reuse: SegmentRun,
) -> None:
    """**Validates: Requirements 2.3**

    共享对象是**完整 DOM**（可写）。8 个消费方之间只要有一个就地改了它，后面的消费方读到
    的就不是文件里的东西 —— 而 `collect_workbook_structure` 恰恰按 `ws._cells` 的成员关系
    判「这个转置字段格存不存在」。

    判据面是 `_cells` **键集合摘要**：同一张受管 sheet 在 4 次共享调用之间必须逐次相同，
    并且段末（最后一个消费方之后）再取一次仍然相同 —— 最后那次覆盖「最后一个消费方改了」
    这个前 4 次盖不到的窗口。
    """
    calls = [call for call in with_reuse.artifact_resolve_calls]
    by_sheet: dict[str, list[ResolveCall]] = {}
    for call in calls:
        by_sheet.setdefault(call.sheet_title, []).append(call)
    assert len(by_sheet) == TRANSPOSED_SPECS, (
        f"共享调用只覆盖了 {len(by_sheet)} 张受管 sheet：{sorted(by_sheet)}"
    )
    for title, sheet_calls in sorted(by_sheet.items()):
        digests = {call.cells for call in sheet_calls}
        assert len(digests) == 1, (
            f"sheet {title!r} 的 `_cells` 键集合在 {len(sheet_calls)} 次共享调用之间变了"
            f"（{len(digests)} 种）⇒ 其中某个消费方就地改了共享 workbook：{sheet_calls}"
        )
        assert with_reuse.cells_after_segment[title] == digests.pop(), (
            f"sheet {title!r} 在段末的 `_cells` 键集合与段中不同 ⇒ 最后一个消费方"
            "（structure_hash 侧）改了共享 workbook"
        )


def test_extract_and_verify_are_still_two_independent_calls() -> None:
    """**Validates: Requirements 2.3**

    design 六 第 3 条：**不**把 extract/verify 合并成一个函数 —— verify 的独立存在就是
    「不信任 materialize 的自述」。本任务只复用解析，判据要能在有人顺手合并时打红：

    * 生产 CPU 段里 `adapter.extract(` 与 `adapter.verify_unmanaged_regions(` 两句都在；
    * `verify_unmanaged_regions` 的形参里**没有** extract 的产物 —— 它自己从字节重算，
      不接受别人递过来的结论；
    * `assert_equivalent()` 仍被调用（verify 的结论仍然是门，不是报告）。
    """
    from app.services.workpaper_sync.adapters.excel import ExcelSyncAdapter

    scoped = inspect.getsource(ContentMutationService._stage_cpu_segment_scoped)
    assert "adapter.extract(" in scoped and "adapter.verify_unmanaged_regions(" in scoped
    assert "unmanaged.assert_equivalent()" in scoped, (
        "CPU 段不再调 `assert_equivalent()` ⇒ verify 从门退化成报告"
    )
    signature = inspect.signature(ExcelSyncAdapter.verify_unmanaged_regions)
    forbidden = {"extracted", "projection", "extract_outcome"}
    leaked = forbidden & set(signature.parameters)
    assert not leaked, (
        f"`verify_unmanaged_regions` 开始接受 extract 侧的产物 {leaked} ⇒ 复用了**结论**，"
        "需求 2.3 与 design 六 第 3 条都禁止"
    )


# ═══════════════════════════════════════════════════════════════════════════
# 五、P4 的 property 形态：解析次数与消费方个数解耦
# ═══════════════════════════════════════════════════════════════════════════


@settings(max_examples=5, deadline=None, suppress_health_check=[HealthCheck.too_slow])
@given(consumers=st.integers(min_value=1, max_value=8))
def test_parse_count_is_decoupled_from_consumer_count(consumers: int) -> None:
    """**Validates: Requirements 2.1**　**Property: P4**

    P4 的命题不是「恰好是 1 次」这个数字，而是**解耦**：同一份字节、同一个视图，不管有
    几个消费方要读，作用域内的解析次数恒为 1。真库那条判据只覆盖 `consumers == 8`
    这一个点；本条把 1..8 都跑一遍，钉住的是「与个数无关」而不是「8 的时候恰好对」。

    同时断另一半：换一个视图（`data_only=True`）必须**再**解析一次（需求 2.2 —— 两个
    视图不可互相冒充），所以总数是「视图数」而不是「1」。
    """
    data = _tiny_workbook_bytes()
    parses: list[tuple[bool, bool]] = []
    import openpyxl

    original = openpyxl.load_workbook

    def counted(*args: Any, **kwargs: Any) -> Any:
        parses.append((bool(kwargs.get("data_only")), bool(kwargs.get("read_only"))))
        return original(*args, **kwargs)

    openpyxl.load_workbook = counted  # type: ignore[assignment]
    try:
        with workbook_read_scope():
            books = [shared_workbook_from_bytes(data) for _ in range(consumers)]
            assert all(book is books[0] for book in books), (
                f"{consumers} 个消费方拿到了不同的对象 ⇒ 没复用"
            )
            assert len(parses) == 1, (
                f"{consumers} 个消费方触发了 {len(parses)} 次解析（应为 1）：{parses}"
            )
            other_view = shared_workbook_from_bytes(data, data_only=True)
            assert other_view is not books[0]
            assert len(parses) == 2, (
                f"换视图后解析 {len(parses)} 次（应为 2）⇒ 两个 data_only 视图互相冒充了"
            )
            assert len(byte_scoped_keys()) == 2
    finally:
        openpyxl.load_workbook = original  # type: ignore[assignment]


# ═══════════════════════════════════════════════════════════════════════════
# 六、同一份字节的**纯 XML/zip 解析**也只做一次（需求 2.1 的另一半）
#
# 需求 2 的用户故事原话：「我不希望『读同一个文件三次』这种成本被当成固有成本接受。」
# 它说的不只是 `openpyxl.load_workbook` —— 真库 D4 的 CPU 段里还有两处「每消费方一遍」：
# `xl/workbook.xml` + 46 张 sheet 的 `<tableParts>`（78 个 binding 各一遍）与
# `xl/sharedStrings.xml` 的前缀 digest（117 次 unmanaged 比对各一遍）。
# ═══════════════════════════════════════════════════════════════════════════


#: 段内出现过的、**有文件身份**的字节形态数：substrate + 转置写入前的中间产物 + 最终产物。
#: 记忆化按 `(路径, mtime_ns, size)` 定身份 ⇒ 每个形态各解析一次是可达下界。
SEGMENT_IDENTIFIED_BYTE_STATES = 3


def test_repeated_zip_xml_parses_collapse_to_one_per_file(
    with_reuse: SegmentRun, world: D4World
) -> None:
    """**Validates: Requirements 2.1**

    三处记忆化各断两件事：**调用次数**（分母 —— 真的有很多消费方在读）与**真解析次数**
    （必须塌到「段内的字节形态数」这个下界）。实测 `_sheet_part_map` 在整段被调 **549**
    次，而有身份的真解析只剩 6 次（3 个字节形态 × workbook_xml/tables 各一次）。

    ⚠️ 诚实分流：`excel_materialize` 的单趟计划阶段在 **BytesIO** 上开 zip 读 sheet part
    映射（每 binding 一次），那种 zip 拿不到文件身份 ⇒ 按 `scoped_parse_memo` 的约定退化
    成不记忆化。这些**不**算记忆化失效，但也不许藏起来：本判据把它们单独列出并按 binding
    数钉死，且只涉及便宜的 `parse_workbook_xml`（读 `xl/workbook.xml` + rels），
    贵的 `parse_tables`（要读 46 张 sheet part）在那条路径上一次都没被调。
    """
    tally = with_reuse.xml_parses
    for label, wrapper, inner in XML_PARSE_PROBES:
        calls = tally[f"{label}_calls"]
        identified = tally[f"{label}_parses_identified"]
        assert calls > identified, (
            f"{wrapper} 被调 {calls} 次、{inner} 在有身份的字节上真解析 {identified} 次"
            " —— 没有出现复用；要么记忆化没接上，要么这一段压根没有多消费方（分母是空的）"
        )
        assert calls >= world.binding_count, (
            f"{wrapper} 只被调了 {calls} 次（binding {world.binding_count}）⇒ "
            "分母太小，本条判据接近空转"
        )
        assert identified <= SEGMENT_IDENTIFIED_BYTE_STATES, (
            f"{inner} 在有身份的字节上解析了 {identified} 次，超过字节形态数 "
            f"{SEGMENT_IDENTIFIED_BYTE_STATES} ⇒ 同一份字节被解析了不止一次"
        )
    # 贵的那个（46 张 sheet part 全读）必须一次 BytesIO 兜底都没有。
    assert tally["tables_parses_anonymous"] == 0, (
        f"Excel Table 清册在无身份的 zip 上解析了 "
        f"{tally['tables_parses_anonymous']} 次 ⇒ 那一份 30ms 的解析没被记忆化覆盖"
    )
    assert tally["shared_strings_parses_anonymous"] == 0
    # 便宜的那个：单趟计划阶段每 binding 一次（BytesIO，无身份）—— 登记在案的剩余杠杆。
    assert tally["workbook_xml_parses_anonymous"] == world.binding_count, (
        f"BytesIO 上的 `xl/workbook.xml` 解析 {tally['workbook_xml_parses_anonymous']} 次 "
        f"≠ binding 数 {world.binding_count} —— 单趟计划阶段的形态变了，"
        "design 附录 F 里登记的剩余杠杆要相应更新"
    )


def test_the_memoised_parse_results_were_never_mutated_in_place(
    with_reuse: SegmentRun,
) -> None:
    """**Validates: Requirements 2.1, 2.3**

    记忆化条目是**共享的可变对象**（`sheets` / `defined_names` / `tables` 都是 list[dict]）。
    有消费方就地改了它，后面的消费方读到的就不是文件里的事实 —— 那正是「复用解析」变成
    「复用一个被改过的结论」的形态。

    oracle 是**段末重新解析一遍**：拿段内留下的记忆化内容与一次干净解析逐字比。相等 ⇒
    整段没人改过它。判据自带分母：记忆化条目必须真的存在（`workbook_xml` / `tables` 两类
    都要有产物那一份）。
    """
    import zipfile as zf_module

    identity = None
    with zf_module.ZipFile(with_reuse.artifact) as archive:
        identity = EE._zip_identity(archive)
        assert identity is not None, "产物 zip 拿不到文件身份 ⇒ 记忆化压根不会发生"
        fresh = {
            "workbook_xml": repr(EE.parse_workbook_xml(archive)),
            "tables": repr(EE.parse_tables(archive, EE.parse_workbook_xml(archive)[0])),
        }
    for kind, fresh_repr in fresh.items():
        key = (EE._PARSE_MEMO_TAG, kind, identity)
        assert key in with_reuse.parse_memo, (
            f"段内没有产物的 `{kind}` 记忆化条目 ⇒ 本条的分母是空的；"
            f"实有条目：{sorted(k[1] for k in with_reuse.parse_memo)}"
        )
        assert with_reuse.parse_memo[key] == fresh_repr, (
            f"产物的 `{kind}` 记忆化内容与干净解析不同 ⇒ 段内有人就地改了共享的解析结果"
        )


def test_parse_memo_entries_share_the_workbook_scope_without_being_closed_as_workbooks() -> (
    None
):
    """**Validates: Requirements 2.2, 1.4**

    记忆化条目与 workbook 条目住在**同一个**作用域 dict 里（需求 2.2：不得新引入第二套
    缓存），但它们不是 workbook：没有句柄、没有 `close`。三条结构事实：

    * 键前缀 `parse-memo` 不可能等于 `Path.resolve()` 的输出 ⇒ 不会被按路径释放误清、
      也不会遮住某个路径条目；
    * 作用域退出时 `_close_if_closeable` 按「有没有 close」分流，不靠
      `except Exception: pass` 把 `AttributeError` 一起吞掉（吞掉的话真的关闭失败也看不见）；
    * 作用域退出后条目一个不剩（没有长存缓存）。
    """
    assert EE._PARSE_MEMO_TAG == "parse-memo"
    calls: list[int] = []

    class _Closeable:
        def close(self) -> None:
            calls.append(1)

    with workbook_read_scope():
        value = EE.scoped_parse_memo("probe", ("id", 1, 2), lambda: ["shared", "list"])
        assert value == ["shared", "list"]
        assert EE.scoped_parse_memo("probe", ("id", 1, 2), lambda: ["other"]) is value, (
            "同一个 key 第二次没命中记忆化"
        )
        keys = [k for k in scoped_keys() if k[0][0] == EE._PARSE_MEMO_TAG]
        assert len(keys) == 1, f"记忆化条目数 {len(keys)}：{keys}"
        # 没有 close 的条目不得让作用域退出时炸，也不得被当成 workbook 关掉。
        cache = EE._workbook_scope.get()
        assert cache is not None
        cache[(("sentinel-closeable",), False)] = _Closeable()
    assert calls == [1], "作用域退出时该关的没关（分流把 workbook 也跳过了）"
    with workbook_read_scope():
        assert [k for k in scoped_keys() if k[0][0] == EE._PARSE_MEMO_TAG] == [], (
            "记忆化条目活过了作用域 ⇒ 有人把它挂到了模块级"
        )
    # 无身份（BytesIO 上打开的 zip）时退化成不记忆化，且不得抛。
    assert EE.scoped_parse_memo("probe", None, lambda: "computed") == "computed"

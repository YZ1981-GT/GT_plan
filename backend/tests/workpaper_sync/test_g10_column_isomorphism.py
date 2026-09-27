# -*- coding: utf-8 -*-
"""G10-2 列同构判据 —— 堵住「猜测映射」（照 G9 同构，单区版）。

spec: `g-cycle-single-region-detail-lanes` · Task 9 / C-7
　　　设计依据 `evidence/task8-c6-remaining-eight-template-logic.md` §2

═══ 为什么必须有这组判据 ═══

既有判据只校验「**声明与模板几何**一致」（行号 / 列字母 / 公式列），
**不校验**「前端字段语义与模板列语义一致」。C-6 实测发现 G10 改造前的前端 35 列里有
**16 列**与权威模板不符（OCI/减值四列属 FVOCI 口径、层次/估值属 G10-5/G10-6、衍生三列
属 G10-8、另两列是与模板口径冲突的自研列）—— 在那种状态下硬凑映射会**通过全部既有
判据**却把 A 列的值写进 B 列。

本组判据四层闭环：
  ① provider 的 `header_text` 逐列 == **模板逐格实测值**（模板是权威）
  ② provider 的 `json_key` 集合/顺序 == **前端行接口字段**（前端与声明不得脱钩）
  ③ 公式列与模板公式逐条一致，且钉住三条负债侧口径
  ④ 单区几何：footer 不受管 / 无分段声明 / 读写不盖 section

🔴 **与 G9 判据的关键差异（照抄会假红/假绿）**：
* 「是否分组列」必须按**该列是否落在 R9 的横向合并区内**判，**不能**按「R10 有叶子」判
  —— G10 的 `F`（期初调整数）/`G`（期初审定数）在 R10 有文本但 R9 **无合并区**
  （负债侧调整与审定都是单列、不拆分量）。照 G9 的「R10 有叶子 ⇒ 必须声明
  group_header_cell」会把这两列判红。
* G10 是**单区** ⇒ 无 `row_section_field`，第④组从「三区过滤」改成「不得声明分段」。
"""
from __future__ import annotations

import os
import re
import sys
from pathlib import Path

import openpyxl
import pytest

_REPO = Path(__file__).resolve().parents[3]
_BACKEND = _REPO / "backend"
if str(_BACKEND) not in sys.path:
    sys.path.insert(0, str(_BACKEND))
os.environ.setdefault("DB_DISABLE_SSL", "True")

from app.services.workpaper_sync.phase5_g10_02_detail import (  # noqa: E402
    FIELD_SPECS_G1002,
    FOOTER_ROW_G1002,
    FORMULA_COLUMNS_G1002,
    FORMULA_TEMPLATES_G1002,
    MANAGED_SHEET_G1002,
    SPEC_G1002,
    STORE_ITEM_ID_G1002,
)
from app.services.workpaper_sync.phase5_row_table_sheet import (  # noqa: E402
    iter_store_rows,
    merge_projection_into_store_rows,
)

TPL = _BACKEND / "wp_templates" / "G" / "G10 交易性金融负债.xlsx"
FE_COMPOSABLE = (
    _REPO
    / "audit-platform/frontend/src/components/workpaper/composables/useG10Detail.ts"
)

#: 前端行接口里**不受管**的字段（身份 / 显示序号）。G10 单区 ⇒ **无** `section`。
_NON_MANAGED_FE_FIELDS = {"rowId", "seq"}


@pytest.fixture(scope="module")
def ws():
    wb = openpyxl.load_workbook(TPL, data_only=False)
    try:
        yield wb[MANAGED_SHEET_G1002]
    finally:
        wb.close()


def _r9_group_starts(ws) -> set[str]:
    """R9 的**横向**合并区起始格（跨两行的单列合并不算分组）。"""
    return {
        rng.coord.split(":")[0]
        for rng in ws.merged_cells.ranges
        if rng.min_row == 9 and rng.max_row == 9 and rng.max_col > rng.min_col
    }


def _group_start_of(ws, col_index: int) -> str:
    """该列所属的 R9 横向分组起始格；不属任何分组则 `""`。"""
    for rng in ws.merged_cells.ranges:
        if (
            rng.min_row == 9
            and rng.max_row == 9
            and rng.max_col > rng.min_col
            and rng.min_col <= col_index <= rng.max_col
        ):
            return rng.coord.split(":")[0]
    return ""


def _frontend_row_fields() -> list[str]:
    """按值抠出 `useG10Detail.G10DetailRow` 的字段名（顺序即声明顺序）。"""
    src = FE_COMPOSABLE.read_text(encoding="utf-8")
    m = re.search(r"export interface G10DetailRow \{(?P<body>.*?)\n\}", src, re.S)
    assert m, "找不到 G10DetailRow 接口 —— 前端结构变了，判据须改写"
    out: list[str] = []
    for line in m.group("body").splitlines():
        fm = re.match(r"\s*(\w+)\??:", line)
        if fm:
            out.append(fm.group(1))
    return out


# ════════════════════════════════════════════════════════════════════════════
# ① header_text 逐列 == 模板逐格实测（模板是权威，FC-5）
# ════════════════════════════════════════════════════════════════════════════
class TestHeaderTextMatchesTemplateCellByCell:
    def test_19_columns_in_excel_order_a_to_s(self) -> None:
        from openpyxl.utils import get_column_letter

        cols = [f[1] for f in FIELD_SPECS_G1002]
        assert cols == [get_column_letter(i) for i in range(1, 20)], (
            f"列序必须是 A..S 连续 19 列，实得 {cols}"
        )

    def test_effective_column_count_is_19_and_t_to_x_are_empty(self, ws) -> None:
        """🔴 `max_column=24` 含空列 —— 有效列只有 19，UUID 列据此取 T 而非 Y。"""
        from openpyxl.utils import get_column_letter

        effective = [
            get_column_letter(c)
            for c in range(1, ws.max_column + 1)
            if any(
                ws.cell(row=r, column=c).value not in (None, "")
                for r in (9, 10, *range(11, FOOTER_ROW_G1002 + 1))
            )
        ]
        assert effective == [get_column_letter(i) for i in range(1, 20)], (
            f"有效内容列实测 {effective}"
        )
        assert ws.max_column == 24, "模板 max_column 变了 ⇒ UUID 列取值依据要重算"
        assert SPEC_G1002.uuid_col == "T", (
            f"UUID 列应取有效列右移一列 T，实得 {SPEC_G1002.uuid_col!r}"
        )

    @pytest.mark.parametrize(
        "field", FIELD_SPECS_G1002, ids=[f[1] for f in FIELD_SPECS_G1002]
    )
    def test_header_text_equals_template_cell(self, field, ws) -> None:
        """两级表头：跨两行合并的单列取 R9，其余取叶子行 R10。"""
        from openpyxl.utils import column_index_from_string

        _key, col, _mode, _vt, _json_key, header_text, _group = field
        ci = column_index_from_string(col)
        leaf = ws.cell(row=10, column=ci).value
        group = ws.cell(row=9, column=ci).value
        expect = (
            str(leaf).strip()
            if (leaf and str(leaf).strip())
            else str(group or "").strip()
        )
        assert header_text == expect, (
            f"{col} 列 header_text 实得 {header_text!r}，模板逐格为 {expect!r}"
            f"（R9={group!r} / R10={leaf!r}）"
        )

    @pytest.mark.parametrize(
        "field", FIELD_SPECS_G1002, ids=[f[1] for f in FIELD_SPECS_G1002]
    )
    def test_group_header_cell_follows_r9_merge_geometry(self, field, ws) -> None:
        """🔴 分组归属按「该列是否落在 R9 **横向**合并区内」判，不按 R10 有无叶子判。

        G10 的 `F`/`G` 在 R10 有文本但 R9 无合并区 —— 照 G9 的口径会把这两列判红。
        """
        from openpyxl.utils import column_index_from_string

        _key, col, _mode, _vt, _json_key, _header, group_cell = field
        expect = _group_start_of(ws, column_index_from_string(col))
        assert group_cell == expect, (
            f"{col} 列声明 group_header_cell={group_cell!r}，"
            f"按 R9 横向合并区实测应为 {expect!r}"
        )

    def test_adjustment_and_audited_columns_are_single_not_split(self, ws) -> None:
        """🔴 负债侧核心差异：期初/期末的调整与审定都是**单列**，R9 无分组。

        资产侧 G9 把调整拆成「成本」「公允价值变动」两列（M/N）、审定同样两列（P/Q）。
        照 G9 的形态给 G10 造两列调整，会与模板 F/N 单列错位一整列。
        """
        from openpyxl.utils import column_index_from_string

        for col, leaf in (
            ("F", "期初调整数"),
            ("G", "期初审定数"),
        ):
            ci = column_index_from_string(col)
            assert ws.cell(row=10, column=ci).value == leaf
            assert _group_start_of(ws, ci) == "", f"{col} 不应属任何 R9 分组"
        # 期末的调整/审定虽在「期末余额」K9:O9 分组内，但各自只占一列
        assert _group_start_of(ws, column_index_from_string("N")) == "K9"
        assert _group_start_of(ws, column_index_from_string("O")) == "K9"
        assert [f[4] for f in FIELD_SPECS_G1002 if f[1] in {"F", "G", "N", "O"}] == [
            "openingAdjustment",
            "openingAdjusted",
            "closingAdjustment",
            "closingAdjusted",
        ]

    def test_group_header_cells_point_at_real_merged_group_starts(self, ws) -> None:
        """每个 `group_header_cell` 必须是模板里真实的一级分组合并区起始格。"""
        declared = {f[6] for f in FIELD_SPECS_G1002 if f[6]}
        starts = _r9_group_starts(ws)
        assert declared == starts, (
            f"声明的分组起始格 {sorted(declared)} 与模板 R9 横向合并区 "
            f"{sorted(starts)} 不一致"
        )
        # 实测恰为三组：期初余额 C9:E9 / 本期变动 H9:J9 / 期末余额 K9:O9
        assert sorted(starts) == ["C9", "H9", "K9"], f"一级分组实得 {sorted(starts)}"


# ════════════════════════════════════════════════════════════════════════════
# ② json_key 集合 == 前端行接口字段集合（声明与前端不得脱钩）
# ════════════════════════════════════════════════════════════════════════════
class TestJsonKeysMatchFrontendRowInterface:
    def test_declared_json_keys_equal_frontend_managed_fields(self) -> None:
        """🔴 本组判据的核心：**双向**相等，不是单向包含。

        - 前端多出字段 ⇒ 有列没被受管（OO 侧编辑会丢）
        - 声明多出字段 ⇒ 指向前端不存在的键（投影恒空）
        """
        declared = {f[4] for f in FIELD_SPECS_G1002}
        frontend = set(_frontend_row_fields()) - _NON_MANAGED_FE_FIELDS
        assert declared == frontend, (
            "声明的 json_key 与前端 G10DetailRow 字段必须**双向相等**。\n"
            f"  前端有而声明缺：{sorted(frontend - declared)}\n"
            f"  声明有而前端缺：{sorted(declared - frontend)}"
        )
        assert len(declared) == 19

    def test_frontend_field_order_matches_excel_column_order(self) -> None:
        """前端字段声明顺序 ≡ 模板列序（A..S）—— 顺序一致才能逐列人工核对。"""
        fe = [f for f in _frontend_row_fields() if f not in _NON_MANAGED_FE_FIELDS]
        declared = [f[4] for f in FIELD_SPECS_G1002]
        assert fe == declared, (
            "前端字段顺序应与模板列序一致（便于逐列核对）。\n"
            f"  前端：{fe}\n  声明：{declared}"
        )

    def test_dropped_legacy_fields_are_not_managed_and_each_has_a_home(self) -> None:
        """16 个移除字段都不得回到受管列，且每条都要指得出归属（指不出的才是真冗余）。"""
        src = FE_COMPOSABLE.read_text(encoding="utf-8")
        m = re.search(
            r"DROPPED_LEGACY_G10_FIELDS[^=]*=\s*\[(?P<body>.*?)\n\]", src, re.S
        )
        assert m, "找不到 DROPPED_LEGACY_G10_FIELDS —— 迁移台账被删了"
        dropped = re.findall(r"field:\s*'(\w+)'", m.group("body"))
        reasons = re.findall(r"reason:\s*'([^']+)'", m.group("body"))
        assert len(dropped) == 16, f"移除字段实得 {len(dropped)} 条：{dropped}"
        assert len(reasons) == len(dropped), "每个移除字段必须给出归属"
        declared = {f[4] for f in FIELD_SPECS_G1002}
        assert not (set(dropped) & declared), (
            f"移除字段又出现在受管列里：{sorted(set(dropped) & declared)}"
        )
        # 方向错的三张表归属必须逐字指名（不是笼统「模板没有」）
        joined = " ".join(reasons)
        for sheet in ("公允价值测试表G10-5", "衍生金融工具核查表G10-8"):
            assert sheet in joined, f"移除理由里未指名权威源 {sheet}"

    def test_row_identity_is_not_a_managed_column_and_no_section(self) -> None:
        """`rowId`/`seq` 不占模板任何一列；G10 单区 ⇒ 前端**不应**有 `section`。"""
        declared = {f[4] for f in FIELD_SPECS_G1002}
        for f in _NON_MANAGED_FE_FIELDS:
            assert f not in declared, f"{f} 不应出现在 field_specs 里"
        assert SPEC_G1002.row_identity_key == "rowId"
        assert "section" not in _frontend_row_fields(), (
            "G10 是单区，前端行接口不应有 section 字段（那是 G9 三区的形态）"
        )


# ════════════════════════════════════════════════════════════════════════════
# ③ 公式列与模板公式逐条一致（含负债侧三条口径）
# ════════════════════════════════════════════════════════════════════════════
class TestFormulaColumnsMatchTemplate:
    def test_formula_columns_are_exactly_the_template_formula_cells(self, ws) -> None:
        from openpyxl.utils import get_column_letter

        actual = tuple(
            get_column_letter(c)
            for c in range(1, ws.max_column + 1)
            if isinstance(ws.cell(row=11, column=c).value, str)
            and str(ws.cell(row=11, column=c).value).startswith("=")
        )
        assert FORMULA_COLUMNS_G1002 == actual, (
            f"公式列声明 {FORMULA_COLUMNS_G1002} 与模板 R11 实测 {actual} 不一致"
        )
        assert len(FORMULA_COLUMNS_G1002) == 6

    @pytest.mark.parametrize("col", sorted(FORMULA_TEMPLATES_G1002))
    def test_formula_template_renders_to_every_data_row(self, col: str, ws) -> None:
        from openpyxl.utils import column_index_from_string

        ci = column_index_from_string(col)
        for row in range(SPEC_G1002.first_data_row, SPEC_G1002.last_data_row + 1):
            got = ws.cell(row=row, column=ci).value
            assert got == FORMULA_TEMPLATES_G1002[col].format(r=row), (
                f"{col}{row} 模板实得 {got!r}，声明渲染为 "
                f"{FORMULA_TEMPLATES_G1002[col].format(r=row)!r}"
            )

    def test_closing_fv_accum_includes_interest_column_j(self) -> None:
        """🔴 `L=D+I+J` **含利息 J** —— 交易性金融负债的利息计入财务费用**同时**增加
        负债账面价值。写成 `D+I`（资产侧的形态）会让审定数系统性偏小。
        """
        assert FORMULA_TEMPLATES_G1002["L"] == "=D{r}+I{r}+J{r}"
        assert "J{r}" in FORMULA_TEMPLATES_G1002["L"]

    def test_closing_initial_amount_rolls_from_unaudited_line(self) -> None:
        """🔴 `K=C+H` 走**未审线** —— 不得从审定数 G 推（同 G9 的 `P=C+M`）。"""
        assert FORMULA_TEMPLATES_G1002["K"] == "=C{r}+H{r}"
        assert "G" not in FORMULA_TEMPLATES_G1002["K"]

    def test_fair_value_identity_repeats_on_both_ends(self) -> None:
        """公允价值 = 初始确认金额 + 累计公允价值变动（期初 E、期末 M 两处）。"""
        assert FORMULA_TEMPLATES_G1002["E"] == "=C{r}+D{r}"
        assert FORMULA_TEMPLATES_G1002["M"] == "=K{r}+L{r}"

    def test_audited_is_fair_value_plus_single_adjustment_column(self) -> None:
        """审定数 = 公允价值 + **单列**调整数（期初 G=E+F、期末 O=M+N）。"""
        assert FORMULA_TEMPLATES_G1002["G"] == "=E{r}+F{r}"
        assert FORMULA_TEMPLATES_G1002["O"] == "=M{r}+N{r}"

    def test_no_decrease_column_exists_in_template(self, ws) -> None:
        """🔴 本期变动是**净额列**（表头逐字「增加"+"/减少"—"」）⇒ 模板无「本期减少」。"""
        group = str(ws["H9"].value or "")
        assert "增加" in group and "减少" in group, f"H9 表头实得 {group!r}"
        leaves = {
            str(ws.cell(row=10, column=c).value or "") for c in range(8, 11)
        }  # H/I/J
        assert not any("本期减少" == t for t in leaves), f"H..J 叶子标题实得 {leaves}"
        declared = {f[4] for f in FIELD_SPECS_G1002}
        assert "currentDecrease" not in declared
        assert "closingBalance" not in declared


# ════════════════════════════════════════════════════════════════════════════
# ④ 单区几何：footer 不受管 / 不声明分段 / 读写不盖 section
# ════════════════════════════════════════════════════════════════════════════
class TestSingleRegionGeometry:
    def test_single_managed_region_r11_to_r20(self) -> None:
        assert (SPEC_G1002.first_data_row, SPEC_G1002.last_data_row) == (11, 20)
        assert SPEC_G1002.footer_row == FOOTER_ROW_G1002 == 21
        assert SPEC_G1002.footer_row > SPEC_G1002.last_data_row, (
            "footer 必须在受管区之外"
        )
        assert SPEC_G1002.store_item_id == STORE_ITEM_ID_G1002 == "G10-detail-rows"

    def test_footer_is_column_wise_sum_with_o21_exception(self, ws) -> None:
        """🔴 footer 逐列 `=SUM(x11:x20)`，**唯一例外** `O21=M21+N21`。

        O 列（期末审定数）在数据行是 `=M+N`，合计行若照抄 SUM 会与 M21+N21 双源；
        模板选择了后者。这条差异是「footer 不受管」的具体理由之一 —— 引擎按统一
        pattern 重写 footer 会把它抹平。
        """
        for col in ("C", "D", "E", "F", "G", "H", "I", "J", "K", "L", "M", "N", "R"):
            assert ws[f"{col}21"].value == f"=SUM({col}11:{col}20)", (
                f"{col}21 实得 {ws[f'{col}21'].value!r}"
            )
        assert ws["O21"].value == "=M21+N21", f"O21 实得 {ws['O21'].value!r}"
        assert ws["A21"].value == SPEC_G1002.footer_marker == "合计"

    def test_no_section_declaration_on_single_region_spec(self) -> None:
        """单区 ⇒ `row_section_field` 必须为空（声明了就会把全表行过滤成空）。"""
        assert not SPEC_G1002.row_section_field, (
            f"G10 是单区却声明了 row_section_field={SPEC_G1002.row_section_field!r}"
        )
        assert not SPEC_G1002.row_section_value

    def test_ghost_row_anchor_points_at_column_b_project_name(self) -> None:
        """🔴 幽灵行锚点指 B 列「项目」而不是默认的 A 列「类别」。

        A 列是枚举且模板 R11 本就有预填值「指定类」—— 用它当锚点会让「只填了类别的
        空行」通不过幽灵行防护、而「OO 侧只填了项目名的真行」被当幽灵行剔除。
        """
        assert SPEC_G1002.ghost_row_anchor_index == 1
        anchor = FIELD_SPECS_G1002[SPEC_G1002.ghost_row_anchor_index]
        assert (anchor[1], anchor[4]) == ("B", "liabilityName"), (
            f"锚点列实得 {anchor[1]}/{anchor[4]}"
        )

    def test_iter_reads_every_row_without_section_filtering(self) -> None:
        payload = [
            {"rowId": "a1", "liabilityName": "甲债券"},
            {"rowId": "a2", "liabilityName": "乙债券"},
            {"rowId": "a3", "liabilityName": "丙债券"},
        ]
        assert [rid for rid, _row in iter_store_rows(SPEC_G1002, payload)] == [
            "a1",
            "a2",
            "a3",
        ]

    def test_duplicate_identity_still_fails_closed(self) -> None:
        from app.services.workpaper_sync.phase5_row_table_sheet import (
            RowTableStorePayloadError,
        )

        dup = [{"rowId": "x"}, {"rowId": "x"}]
        with pytest.raises(RowTableStorePayloadError, match="重复行身份"):
            list(iter_store_rows(SPEC_G1002, dup))

    def test_merge_does_not_stamp_section_on_new_rows(self) -> None:
        """OO 侧插的新行**不**带 `section`（那是 G9 三区才需要的盖章）。"""

        class _FV:
            def __init__(self, key: str, row_key: str, value):
                self.stable_key = key
                self.row_key = row_key
                self.value = value
                self.is_protected = False

        new_key = f"{SPEC_G1002.table_key}/newrow/liability_name"

        class _Proj:
            def stable_keys(self):
                return [new_key]

            def get(self, key):
                return _FV(key, "newrow", "OO侧新增") if key == new_key else None

        base = [{"rowId": "a1", "liabilityName": "甲债券"}]
        rows, applied, _visited, touched = merge_projection_into_store_rows(
            SPEC_G1002, projection=_Proj(), base_rows=base
        )
        assert applied == 1 and touched == {"newrow"}
        added = next(r for r in rows if r["rowId"] == "newrow")
        assert added["liabilityName"] == "OO侧新增"
        assert "section" not in added, f"单区不应盖 section，实得 {added!r}"
        assert {r["rowId"] for r in rows} == {"a1", "newrow"}

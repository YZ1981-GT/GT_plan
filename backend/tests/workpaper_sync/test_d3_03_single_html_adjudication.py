"""判据：D3-3 `single_html` 裁决的四条依据现在仍然成立 + Property 11 受管契约缺席守卫。

spec: d3-sync-coverage-via-row-table-engine · Task 15 · Requirements 5.1–5.5
裁决: .kiro/specs/d3-sync-coverage-via-row-table-engine/evidence/task15-d3-03-single-html.json

D3-3「预收账款调整分录汇总」是已有专用集中登记同步链的 hub 表（宿主接
`useAdjustmentCentralSync`→后端 `AdjustmentSyncService`），模板无行身份列、借贷平衡 HTML+
后端双门强制、数据区空白待填。据此判 `single_html`（与已判 single_html 的 D4-4 一致），
**不接行表引擎单元格双向**。

本文件把裁决的四条依据钉成测试（防证据腐烂）——任一条不再成立即必红并提示回去重新裁决——
并落 **Property 11**（上游诚实边界红线）守卫：断言 D3-3 **没有**被接成受管 provider，且
变异（把 `D3-aje-rows` 塞进受管契约）⇒ 缺席守卫**必红**。
"""
from __future__ import annotations

import io
import json
import os
import sys
import warnings
from pathlib import Path

import pytest

_REPO = Path(__file__).resolve().parents[3]
_BACKEND = _REPO / "backend"
if str(_BACKEND) not in sys.path:  # pragma: no cover
    sys.path.insert(0, str(_BACKEND))
os.environ.setdefault("DB_DISABLE_SSL", "True")
warnings.filterwarnings("ignore")

from app.services.workpaper_sync import phase5_d3_expansion as EXP  # noqa: E402
from app.services.workpaper_sync import phase5_d3_prepaid_receipts as D3  # noqa: E402

MANAGED_SHEET_D303 = "调整分录汇总表D3-3"
STORE_KEY_D303 = "D3-aje-rows"

_FRONTEND = _REPO / "audit-platform" / "frontend" / "src"
_VERDICT_PATH = (
    _REPO
    / ".kiro/specs/d3-sync-coverage-via-row-table-engine/evidence"
    / "task15-d3-03-single-html.json"
)
_D44_VERDICT_PATH = (
    _REPO
    / ".kiro/specs/d-cycle-sheet-bidirectional-expansion/evidence"
    / "T08-d44-single-html-adjudication.json"
)


@pytest.fixture(scope="module")
def verdict() -> dict:
    return json.loads(_VERDICT_PATH.read_text(encoding="utf-8"))


@pytest.fixture(scope="module")
def d44_verdict() -> dict:
    return json.loads(_D44_VERDICT_PATH.read_text(encoding="utf-8"))


@pytest.fixture(scope="module")
def sheet():
    import openpyxl

    wb = openpyxl.load_workbook(
        io.BytesIO(D3.read_authoritative_template()), data_only=False
    )
    return wb[MANAGED_SHEET_D303]


# ═══════════════════════════════════════════════════════════════════════════════
# 裁决未腐烂
# ═══════════════════════════════════════════════════════════════════════════════
class TestDecisionIsStillSingleHtml:
    def test_verdict_file_says_single_html(self, verdict: dict) -> None:
        assert verdict["decision"] == "single_html"
        assert len(verdict["why_single_html"]) == 4, "四条依据缺一即裁决不完整"

    def test_no_code_changed_flag(self, verdict: dict) -> None:
        assert verdict["no_code_changed"] is True, (
            "Task 15 是可行性核，不改任何生产代码——裁决文件必须记 no_code_changed=True"
        )


# ═══════════════════════════════════════════════════════════════════════════════
# 依据①：无行身份列
# ═══════════════════════════════════════════════════════════════════════════════
class TestCriterion1NoRowIdentityColumn:
    def test_no_uuid_or_gtrow_trace_anywhere(self, sheet) -> None:
        hits = [
            (cell.coordinate, cell.value)
            for row in sheet.iter_rows()
            for cell in row
            if isinstance(cell.value, str)
            and any(t in cell.value.upper() for t in ("GTROW", "_GT", "UUID", "ROWID"))
        ]
        assert not hits, (
            f"D3-3 模板出现行身份列痕迹 {hits} —— 若确实注入了 UUID 列，"
            "single_html 裁决的依据①不再成立，须回去重新裁决"
        )

    def test_d3_2_has_identity_column_as_contrast(self) -> None:
        """反面对照：真正受管的 D3-2 **有** `uuid_col`，证明本判据能区分两者。"""
        spec = D3.instrumentation_spec()
        assert str(spec.managed_sheet) != MANAGED_SHEET_D303
        assert str(spec.uuid_col).strip(), (
            "D3-2 应有非空 uuid_col 作行身份锚——若它也空了，说明本判据的对照前提失效"
        )


# ═══════════════════════════════════════════════════════════════════════════════
# 依据②：已有专用集中登记同步链
# ═══════════════════════════════════════════════════════════════════════════════
class TestCriterion2DedicatedSyncChainExists:
    def test_d3_tab_adjustment_wires_central_sync_once(self) -> None:
        """接线点：恰 1 条 import + 1 处 useAdjustmentCentralSync({...}) 调用。

        🔴 不断言裸字符串出现次数（实测 = 3，因为 import 行同时写符号名与模块路径）——
        直接断言 3 会把「接线点数」与「字符出现数」混为一谈，下次改 import 写法就误红。
        （范式照 D1-5 判据教训。）
        """
        src = (
            _FRONTEND / "components/workpaper/d3/D3TabAdjustment.vue"
        ).read_text(encoding="utf-8")
        import_lines = [
            ln
            for ln in src.splitlines()
            if ln.strip().startswith("import") and "useAdjustmentCentralSync" in ln
        ]
        assert len(import_lines) == 1, (
            f"预期恰 1 条 import 接线，实得 {len(import_lines)} 条: {import_lines}"
        )
        assert src.count("useAdjustmentCentralSync({") == 1, (
            "预期恰 1 处 useAdjustmentCentralSync({...}) 调用接线——接线点数变了，依据②须复核"
        )
        assert f"itemId: '{STORE_KEY_D303}'" in src
        assert "wpCode: 'D3'" in src

    def test_store_key_is_owned_by_use_d3_adjustment(self) -> None:
        src = (
            _FRONTEND / "components/workpaper/composables/useD3Adjustment.ts"
        ).read_text(encoding="utf-8")
        assert f"const ITEM_ID_ROWS = '{STORE_KEY_D303}'" in src

    def test_progress_key_reads_the_same_store(self) -> None:
        src = (_FRONTEND / "components/workpaper/d3/D3TabIndex.vue").read_text(
            encoding="utf-8"
        )
        assert f"'{STORE_KEY_D303}'" in src, (
            "D3TabIndex 的完成度判定不再读 D3-aje-rows——hub store 的消费面变了，依据②须复核"
        )

    def test_backend_sync_service_validates_balance(self) -> None:
        """后端 AdjustmentSyncService 也是这条链的一环，且它硬校验借贷平衡。"""
        svc_src = (
            _BACKEND / "app/services/adjustment_sync_service.py"
        ).read_text(encoding="utf-8")
        assert "UNBALANCED" in svc_src
        assert "total_debit != total_credit" in svc_src


# ═══════════════════════════════════════════════════════════════════════════════
# 依据③：借贷平衡 HTML+后端双层强制，Excel 模板零校验
# ═══════════════════════════════════════════════════════════════════════════════
class TestCriterion3BalanceEnforcedOffExcel:
    def test_html_side_has_balance_computed_and_gate(self) -> None:
        composable = (
            _FRONTEND / "components/workpaper/composables/useD3Adjustment.ts"
        ).read_text(encoding="utf-8")
        assert "isBalanced" in composable
        assert "balanceDiff" in composable
        host = (
            _FRONTEND / "components/workpaper/d3/D3TabAdjustment.vue"
        ).read_text(encoding="utf-8")
        assert "!isBalanced" in host, (
            "「同步到集中登记」按钮不再以 isBalanced 为门——依据③须复核"
        )

    def test_excel_data_band_has_no_formula_and_no_content(self, sheet) -> None:
        """数据区 6-20 是空模板带：零内容零公式 ⇒ Excel 侧不可能校验平衡。"""
        non_empty = [
            (sheet.cell(row=r, column=c).coordinate, sheet.cell(row=r, column=c).value)
            for r in range(6, 21)
            for c in range(1, 11)
            if sheet.cell(row=r, column=c).value is not None
        ]
        assert not non_empty, (
            f"D3-3 数据区 6-20 出现内容/公式 {non_empty[:5]} —— "
            "若模板加了平衡校验公式，依据③不再成立"
        )

    def test_all_formulas_are_header_directory_refs_not_data_rows(self, sheet) -> None:
        """6 处公式全部是页眉/索引区（行3-4）对底稿目录的引用，无一落在数据行 6-20。"""
        formula_rows = {
            cell.row
            for row in sheet.iter_rows()
            for cell in row
            if isinstance(cell.value, str) and cell.value.startswith("=")
        }
        assert formula_rows, "模板应有页眉引用公式（若一个都没有说明模板变了）"
        assert not (formula_rows & set(range(6, 21))), (
            f"数据行 6-20 出现公式（行 {sorted(formula_rows & set(range(6, 21)))}）——依据③须复核"
        )


# ═══════════════════════════════════════════════════════════════════════════════
# 与 D4-4 同一套模板范式
# ═══════════════════════════════════════════════════════════════════════════════
class TestSamePatternAsD44Precedent:
    def test_headers_match_d4_4_verdict_verbatim(self, sheet, verdict, d44_verdict) -> None:
        actual = [sheet.cell(row=5, column=c).value for c in range(1, 11)]
        assert actual == d44_verdict["authority"]["header_A_to_J"], (
            f"D3-3 表头与 D4-4 不再逐字相同：\nD3-3={actual}\n"
            f"D4-4={d44_verdict['authority']['header_A_to_J']}\n"
            "「同一套模板范式」这条类推依据须复核"
        )
        # 本 spec 自己的裁决文件也必须记着同一份表头（两处不得分叉）。
        assert actual == verdict["authority"]["header_A_to_J"]

    def test_both_verdicts_agree_on_single_html(self, verdict, d44_verdict) -> None:
        assert d44_verdict["decision"] == verdict["decision"] == "single_html"


# ═══════════════════════════════════════════════════════════════════════════════
# Property 11：D3-3 未接成受管 provider（受管契约缺席守卫）+ 变异必红
# ═══════════════════════════════════════════════════════════════════════════════
class TestProperty11D303NotRegisteredInManagedContract:
    """核阶段不改任何生产代码 ⇒ D3-3 不得出现在受管契约里。"""

    def test_store_key_absent_from_all_store_item_ids(self) -> None:
        items = EXP.all_store_item_ids()
        assert STORE_KEY_D303 not in items, (
            f"{STORE_KEY_D303!r} 出现在 all_store_item_ids()={items}——"
            "D3-3 被偷偷接进行表受管契约，违反 single_html 裁决 + Property 11 诚实边界"
        )

    def test_managed_sheet_absent_from_all_managed_sheet_names(self) -> None:
        names = EXP.all_managed_sheet_names()
        assert MANAGED_SHEET_D303 not in names, (
            f"{MANAGED_SHEET_D303!r} 出现在 all_managed_sheet_names()={names}——"
            "D3-3 被接进受管 sheet 集合，违反 single_html 裁决"
        )

    def test_no_spec_declares_d3_03_managed_sheet(self) -> None:
        """任何行表 spec / 审定表 spec 都不得声明 managed_sheet='调整分录汇总表D3-3'。"""
        row_specs = EXP.managed_row_table_specs()
        for s in row_specs:
            assert str(s.managed_sheet) != MANAGED_SHEET_D303, (
                f"行表 spec {s!r} 声明了 D3-3——不该有 provider"
            )
        adj = EXP.adjudication_spec()
        if adj is not None:
            assert str(adj.managed_sheet) != MANAGED_SHEET_D303

    def test_no_phase5_d3_03_module_exists(self) -> None:
        """不存在 phase5_d3_03 provider 模块（对照 06/04/05/07 都存在）。"""
        import importlib

        with pytest.raises(ModuleNotFoundError):
            importlib.import_module("app.services.workpaper_sync.phase5_d3_03_adjustment")


class TestMutationRegisteringD303WouldGoRed:
    """🔴 变异检验：模拟把 D3-3 接进受管契约 ⇒ 缺席守卫必红（证明守卫有牙齿）。"""

    def test_mutation_injecting_store_key_makes_absence_guard_red(self) -> None:
        """把 D3-aje-rows 塞进 store item 集合后，缺席断言必须抛 AssertionError。"""
        real_items = list(EXP.all_store_item_ids())
        mutated_items = tuple(real_items + [STORE_KEY_D303])  # 变异：接进受管契约

        # 对照：真集合缺席断言应通过（不抛）
        assert STORE_KEY_D303 not in real_items

        # 变异体：同一条缺席断言必须抛 —— 守卫能检出「被接入」
        with pytest.raises(AssertionError):
            assert STORE_KEY_D303 not in mutated_items, (
                "变异后 D3-aje-rows 已在受管契约里"
            )

    def test_mutation_injecting_managed_sheet_makes_absence_guard_red(self) -> None:
        real_names = list(EXP.all_managed_sheet_names())
        mutated_names = tuple(real_names + [MANAGED_SHEET_D303])

        assert MANAGED_SHEET_D303 not in real_names
        with pytest.raises(AssertionError):
            assert MANAGED_SHEET_D303 not in mutated_names

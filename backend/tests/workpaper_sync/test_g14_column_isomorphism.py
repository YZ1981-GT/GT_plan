# -*- coding: utf-8 -*-
"""G14-2 列同构判据 —— 固定行集 + 布尔校验列 + 模板缺陷台账。

spec: `g-cycle-single-region-detail-lanes` · Task 10 / C-9
　　　设计依据 `evidence/task8-c6-remaining-eight-template-logic.md` §4

═══ 五层闭环 ═══
  ① `header_text` 逐列 == **模板逐格实测**
  ② `json_key` 集合/顺序 == **前端行接口前 13 字段**（双向相等）
  ③ 🔴 固定行集：provider 与前端的行名表双向锁，且**恰 9 行**
  ④ 🔴 `K=G+H` 模板缺陷台账 + `L` 布尔校验列（裁决 G1R-H4）
  ⑤ 单区几何 / footer L20 例外 / `rowKey` 行身份
"""
from __future__ import annotations

import os
import re
import sys
from pathlib import Path

import openpyxl
import pytest

_REPO = Path(__file__).resolve().parents[3]
_BACKEND = _REPO / "backend"
if str(_BACKEND) not in sys.path:
    sys.path.insert(0, str(_BACKEND))
os.environ.setdefault("DB_DISABLE_SSL", "True")

from app.services.workpaper_sync.phase5_g14_02_detail import (  # noqa: E402
    BOOLEAN_COLUMNS_G1402,
    FIELD_SPECS_G1402,
    FOOTER_ROW_G1402,
    FORMULA_COLUMNS_G1402,
    FORMULA_TEMPLATES_G1402,
    FRONTEND_ONLY_FIELDS_G1402,
    MANAGED_SHEET_G1402,
    SPEC_G1402,
    STORE_ITEM_ID_G1402,
    TEMPLATE_ROW_DEFECTS_G1402,
    TEMPLATE_ROW_LABELS_G1402,
)
from app.services.workpaper_sync.phase5_row_table_sheet import (  # noqa: E402
    iter_store_rows,
    merge_projection_into_store_rows,
)

TPL = _BACKEND / "wp_templates" / "G" / "G14 信用减值损失.xlsx"
_FE = _REPO / "audit-platform/frontend/src/components/workpaper/composables"
FE_COMPOSABLE = _FE / "useG14Detail.ts"
FE_CONSTANTS = _FE / "g14Constants.ts"

#: 前端行接口里**不受管**的字段：行身份 + 三个试算对账派生值
_NON_MANAGED_FE_FIELDS = {"rowKey", *FRONTEND_ONLY_FIELDS_G1402}

_DATA_ROWS = tuple(range(SPEC_G1402.first_data_row, SPEC_G1402.last_data_row + 1))


@pytest.fixture(scope="module")
def ws():
    wb = openpyxl.load_workbook(TPL, data_only=False)
    try:
        yield wb[MANAGED_SHEET_G1402]
    finally:
        wb.close()


def _group_start_of(ws, col_index: int) -> str:
    for rng in ws.merged_cells.ranges:
        if (
            rng.min_row == 9
            and rng.max_row == 9
            and rng.max_col > rng.min_col
            and rng.min_col <= col_index <= rng.max_col
        ):
            return rng.coord.split(":")[0]
    return ""


def _frontend_row_fields() -> list[str]:
    src = FE_COMPOSABLE.read_text(encoding="utf-8")
    m = re.search(r"export interface G14DetailRow \{(?P<body>.*?)\n\}", src, re.S)
    assert m, "找不到 G14DetailRow 接口 —— 前端结构变了，判据须改写"
    return [
        fm.group(1)
        for line in m.group("body").splitlines()
        if (fm := re.match(r"\s*(\w+)\??:", line))
    ]


def _frontend_template_row_labels() -> list[str]:
    src = FE_CONSTANTS.read_text(encoding="utf-8")
    m = re.search(
        r"G14_TEMPLATE_ROW_LABELS\s*=\s*\[(?P<body>.*?)\n\]", src, re.S
    )
    assert m, "找不到 G14_TEMPLATE_ROW_LABELS —— 前端行名真源被删了"
    return re.findall(r"'([^']+)'", m.group("body"))


def _frontend_line_item_labels() -> list[str]:
    src = FE_CONSTANTS.read_text(encoding="utf-8")
    m = re.search(
        r"export const G14_LINE_ITEMS[^=]*=\s*\[(?P<body>[\s\S]*?)\n\]\n", src
    )
    assert m, "找不到 G14_LINE_ITEMS"
    return re.findall(r"^\s*label:\s*'([^']+)'", m.group("body"), re.M)


# ════════════════════════════════════════════════════════════════════════════
# ① header_text 逐列 == 模板逐格实测
# ════════════════════════════════════════════════════════════════════════════
class TestHeaderTextMatchesTemplateCellByCell:
    def test_13_columns_in_excel_order_a_to_m(self) -> None:
        from openpyxl.utils import get_column_letter

        cols = [f[1] for f in FIELD_SPECS_G1402]
        assert cols == [get_column_letter(i) for i in range(1, 14)], (
            f"列序必须是 A..M 连续 13 列，实得 {cols}"
        )

    def test_effective_columns_equal_max_column_so_uuid_is_n(self, ws) -> None:
        """🔴 G14 没有空尾列：有效列 13 **恰等于** `max_column` ⇒ 两个 uuid 口径重合。"""
        from openpyxl.utils import get_column_letter

        effective = [
            get_column_letter(c)
            for c in range(1, ws.max_column + 1)
            if any(
                ws.cell(row=r, column=c).value not in (None, "")
                for r in (9, 10, *_DATA_ROWS, FOOTER_ROW_G1402)
            )
        ]
        assert effective == [get_column_letter(i) for i in range(1, 14)]
        assert ws.max_column == 13, "模板 max_column 变了 ⇒ uuid 列取值依据要重算"
        assert SPEC_G1402.uuid_col == "N"

    @pytest.mark.parametrize(
        "field", FIELD_SPECS_G1402, ids=[f[1] for f in FIELD_SPECS_G1402]
    )
    def test_header_text_equals_template_cell(self, field, ws) -> None:
        from openpyxl.utils import column_index_from_string

        _key, col, _mode, _vt, _json_key, header_text, _group = field
        ci = column_index_from_string(col)
        leaf = ws.cell(row=10, column=ci).value
        group = ws.cell(row=9, column=ci).value
        expect = (
            str(leaf).strip()
            if (leaf and str(leaf).strip())
            else str(group or "").strip()
        )
        assert header_text == expect, (
            f"{col} 列 header_text 实得 {header_text!r}，模板逐格为 {expect!r}"
            f"（R9={group!r} / R10={leaf!r}）"
        )

    @pytest.mark.parametrize(
        "field", FIELD_SPECS_G1402, ids=[f[1] for f in FIELD_SPECS_G1402]
    )
    def test_group_header_cell_follows_r9_merge_geometry(self, field, ws) -> None:
        from openpyxl.utils import column_index_from_string

        _key, col, _mode, _vt, _json_key, _header, group_cell = field
        assert group_cell == _group_start_of(ws, column_index_from_string(col))

    def test_exactly_two_r9_horizontal_groups(self, ws) -> None:
        starts = {
            rng.coord.split(":")[0]
            for rng in ws.merged_cells.ranges
            if rng.min_row == 9 and rng.max_row == 9 and rng.max_col > rng.min_col
        }
        assert sorted(starts) == ["B9", "F9"], f"R9 横向分组实得 {sorted(starts)}"
        assert {f[6] for f in FIELD_SPECS_G1402 if f[6]} == starts


# ════════════════════════════════════════════════════════════════════════════
# ② json_key == 前端行接口（前 13 字段）
# ════════════════════════════════════════════════════════════════════════════
class TestJsonKeysMatchFrontendRowInterface:
    def test_declared_json_keys_equal_frontend_managed_fields(self) -> None:
        declared = {f[4] for f in FIELD_SPECS_G1402}
        frontend = set(_frontend_row_fields()) - _NON_MANAGED_FE_FIELDS
        assert declared == frontend, (
            "声明的 json_key 与前端 G14DetailRow 受管字段必须**双向相等**。\n"
            f"  前端有而声明缺：{sorted(frontend - declared)}\n"
            f"  声明有而前端缺：{sorted(declared - frontend)}"
        )
        assert len(declared) == 13

    def test_frontend_field_order_matches_excel_column_order(self) -> None:
        fe = [f for f in _frontend_row_fields() if f not in _NON_MANAGED_FE_FIELDS]
        assert fe == [f[4] for f in FIELD_SPECS_G1402], (
            f"前端字段顺序应与模板列序一致。\n  前端：{fe}\n  声明：{[f[4] for f in FIELD_SPECS_G1402]}"
        )

    def test_tb_reconciliation_fields_are_frontend_only(self) -> None:
        """三个试算对账派生字段必须在前端接口里、但**不**在受管列里。"""
        declared = {f[4] for f in FIELD_SPECS_G1402}
        fe = set(_frontend_row_fields())
        for f in FRONTEND_ONLY_FIELDS_G1402:
            assert f in fe, f"{f} 不在前端接口里 ⇒ 登记表过期"
            assert f not in declared, f"{f} 不是模板列，不应受管"

    def test_dropped_legacy_fields_are_not_managed_and_each_has_a_home(self) -> None:
        src = FE_COMPOSABLE.read_text(encoding="utf-8")
        m = re.search(
            r"DROPPED_LEGACY_G14_FIELDS[^=]*=\s*\[(?P<body>[\s\S]*?)\n\]", src
        )
        assert m, "找不到 DROPPED_LEGACY_G14_FIELDS —— 迁移台账被删了"
        dropped = re.findall(r"field:\s*'(\w+)'", m.group("body"))
        assert len(dropped) == 4, f"移除字段实得 {len(dropped)} 条：{dropped}"
        declared = {f[4] for f in FIELD_SPECS_G1402}
        assert not (set(dropped) & declared)
        # 🔴 `otherMovement` 的移除理由必须点明「模板 J 不含其他变动项」
        body = m.group("body")
        assert "otherMovement" in dropped
        assert "F+G-H-I" in body or "不含" in body


# ════════════════════════════════════════════════════════════════════════════
# ③ 🔴 固定行集：provider ↔ 前端 ↔ 模板 三方锁
# ════════════════════════════════════════════════════════════════════════════
class TestFixedRowSetIsLockedToTemplate:
    def test_provider_row_labels_equal_template_column_a(self, ws) -> None:
        actual = tuple(str(ws.cell(row=r, column=1).value or "").strip() for r in _DATA_ROWS)
        assert TEMPLATE_ROW_LABELS_G1402 == actual, (
            f"provider 行名表与模板 A 列不一致。\n  声明：{TEMPLATE_ROW_LABELS_G1402}\n  模板：{actual}"
        )
        assert len(TEMPLATE_ROW_LABELS_G1402) == 9

    def test_frontend_template_row_labels_equal_provider(self) -> None:
        """前端 `G14_TEMPLATE_ROW_LABELS` 与 provider 的同名表**双向相等**。"""
        assert _frontend_template_row_labels() == list(TEMPLATE_ROW_LABELS_G1402)

    def test_frontend_line_items_equal_template_row_set(self) -> None:
        """🔴 `G14_LINE_ITEMS` 的 label 序列必须逐项等于模板行集。

        改造前它有 **10** 项（自研第 10 行「合同资产减值损失」）—— 行表引擎按数组顺序
        映射 R11-R19，10 行落进 9 行区会扩行、把 footer R20 挤下去。
        """
        labels = _frontend_line_item_labels()
        assert labels == list(TEMPLATE_ROW_LABELS_G1402), (
            f"G14_LINE_ITEMS 的 label 序列与模板行集不一致：\n  前端 {labels}"
        )
        assert "合同资产减值损失" not in labels

    def test_contract_asset_ecl_is_folded_into_the_other_row(self) -> None:
        """合同资产的三处落点都改指「其他」行 —— 口径不丢、链不断。"""
        constants = FE_CONSTANTS.read_text(encoding="utf-8")
        # ① 1142 试算取数并入 other 行
        m = re.search(
            r"rowKey: 'other',[\s\S]{0,400}?tbPrefixes: \[([^\]]*)\]", constants
        )
        assert m and "1142" in m.group(1), "other 行未收 1142 取数前缀"
        # ② D6 的 ECL 交叉索引挂 other
        assert re.search(r"other: 'wp:D6-1'", constants), "G14_ECL_CROSS_REF.other 未指 D6"
        assert not re.search(r"^\s*ca: 'wp:D6-1'", constants, re.M), "残留 ca 的 ECL 索引"
        # ③ D6 的 ECL 事件落点
        ecl = (_FE / "gCycleSourceEcl.ts").read_text(encoding="utf-8")
        assert re.search(r"D6: 'other'", ecl), "gCycleSourceEcl 的 D6 未改指 other"
        # ④ 含「合同资产」的调整分录 rowKey 推断
        adj = (_FE / "g14AdjStorage.ts").read_text(encoding="utf-8")
        assert re.search(r"/合同资产/\.test\(text\)\) return 'other'", adj)

    def test_dropped_row_ledger_records_why(self) -> None:
        src = FE_COMPOSABLE.read_text(encoding="utf-8")
        m = re.search(r"DROPPED_LEGACY_G14_ROWS[^=]*=\s*\[(?P<body>[\s\S]*?)\n\]", src)
        assert m, "找不到 DROPPED_LEGACY_G14_ROWS"
        assert "ca" in m.group("body")
        assert "其他" in m.group("body"), "移除理由未写明并入哪一行"


# ════════════════════════════════════════════════════════════════════════════
# ④ 🔴 K 列模板缺陷 + L 布尔校验列
# ════════════════════════════════════════════════════════════════════════════
class TestFormulaColumnsAndTemplateDefect:
    def test_formula_columns_are_exactly_the_template_formula_cells(self, ws) -> None:
        from openpyxl.utils import get_column_letter

        every_row = tuple(
            get_column_letter(c)
            for c in range(1, 14)
            if all(
                isinstance(ws.cell(row=r, column=c).value, str)
                and str(ws.cell(row=r, column=c).value).startswith("=")
                for r in _DATA_ROWS
            )
        )
        assert FORMULA_COLUMNS_G1402 == every_row, (
            f"公式列声明 {FORMULA_COLUMNS_G1402} 与「每行都有公式」的实测 {every_row} 不一致"
        )
        assert len(FORMULA_COLUMNS_G1402) == 4

    @pytest.mark.parametrize("col", sorted(FORMULA_TEMPLATES_G1402))
    def test_formula_template_renders_to_every_data_row(self, col: str, ws) -> None:
        """🔴 `formula_templates` **逐字记模板原式**（含 K 的缺陷）⇒ 逐格可比。"""
        from openpyxl.utils import column_index_from_string

        ci = column_index_from_string(col)
        for r in _DATA_ROWS:
            got = ws.cell(row=r, column=ci).value
            assert got == FORMULA_TEMPLATES_G1402[col].format(r=r), (
                f"{col}{r} 模板实得 {got!r}，声明渲染为 "
                f"{FORMULA_TEMPLATES_G1402[col].format(r=r)!r}"
            )

    def test_rollforward_formula_requires_positive_reversal(self) -> None:
        """`J=F+G-H-I` 的 `-H` 是「转回填正数」的依据 —— 它是对的那一式。"""
        assert FORMULA_TEMPLATES_G1402["J"] == "=F{r}+G{r}-H{r}-I{r}"

    def test_k_column_template_defect_is_recorded_not_silently_fixed(self) -> None:
        """🔴 模板 `K=G+H` 与同表 `J` 的 `-H` 矛盾 ⇒ 缺陷台账记下来，声明**不许**改成 `=G-H`。

        改成 `=G-H` 会让 `test_formula_template_renders_to_every_data_row` 必红
        （判据逐格比对模板）。会计正确口径由**前端** `calcNetImpairmentLoss` 承担。
        """
        assert FORMULA_TEMPLATES_G1402["K"] == "=G{r}+H{r}"
        by_col = {c: (rows, note) for c, rows, note in TEMPLATE_ROW_DEFECTS_G1402}
        assert sorted(by_col) == ["K"], f"缺陷台账应只有 K 一列，实得 {sorted(by_col)}"
        rows, note = by_col["K"]
        assert rows == _DATA_ROWS, "K 的缺陷覆盖全部数据行"
        assert "=G{r}-H{r}" in note, "台账须写出会计正确口径"
        assert "转回" in note

    def test_k_defect_still_present_in_template_cell_by_cell(self, ws) -> None:
        """逐格复核缺陷仍在 —— 模板一旦被修成 `=G-H`，这里会红并提示同步更新台账。"""
        for r in _DATA_ROWS:
            got = str(ws[f"K{r}"].value or "")
            assert got == f"=G{r}+H{r}", f"K{r} 实得 {got!r}（模板被修过？请同步台账）"

    def test_frontend_computes_profit_loss_as_provision_minus_reversal(self) -> None:
        """前端按会计正确口径算 —— `calcNetImpairmentLoss` 必须是「计提 − 转回」。"""
        engine = (_FE / "useG14FormulaEngine.ts").read_text(encoding="utf-8")
        m = re.search(
            r"export function calcNetImpairmentLoss\([\s\S]{0,300}?\n\}", engine
        )
        assert m, "找不到 calcNetImpairmentLoss"
        body = m.group(0)
        assert "-" in body and "+" not in body.split("return")[-1], (
            f"calcNetImpairmentLoss 不是减法口径：{body}"
        )

    def test_l_is_a_boolean_check_column(self) -> None:
        """🔴 裁决 G1R-H4：`L=D=K` 求值 TRUE/FALSE ⇒ boolean + formula（不入 store）。"""
        by_col = {f[1]: f for f in FIELD_SPECS_G1402}
        assert by_col["L"][2] == "formula", "布尔列标 editable 会让 TRUE/FALSE 进 store"
        assert by_col["L"][3] == "boolean", "布尔列的 value_type 必须是 boolean"
        assert BOOLEAN_COLUMNS_G1402 == ("L",)
        assert FORMULA_TEMPLATES_G1402["L"] == "=D{r}=K{r}"
        assert by_col["L"][4] == "reconciled"

    def test_only_l_is_boolean_among_managed_fields(self) -> None:
        assert [f[1] for f in FIELD_SPECS_G1402 if f[3] == "boolean"] == ["L"]

    def test_every_formula_mode_field_is_covered_by_formula_mask(self) -> None:
        """契约层 CS-13 的本地复核（提前于 parse_contract 打红，定位更快）。"""
        from app.services.workpaper_sync.contracts import column_in_ranges

        mask = SPEC_G1402.formula_mask
        for key, col, mode, _vt, _json, _hdr, _grp in FIELD_SPECS_G1402:
            if mode != "formula":
                continue
            assert column_in_ranges(col, mask), (
                f"{key}（列 {col}）声明 mode=formula 但不在 formula_mask {mask} 内"
            )


# ════════════════════════════════════════════════════════════════════════════
# ⑤ 几何与 store 读写
# ════════════════════════════════════════════════════════════════════════════
class TestGeometryAndStore:
    def test_fixed_region_r11_to_r19_with_footer_outside(self) -> None:
        assert (SPEC_G1402.first_data_row, SPEC_G1402.last_data_row) == (11, 19)
        assert SPEC_G1402.footer_row == FOOTER_ROW_G1402 == 20
        assert SPEC_G1402.footer_row > SPEC_G1402.last_data_row
        assert SPEC_G1402.store_item_id == STORE_ITEM_ID_G1402 == "G14-detail-rows"

    def test_row_identity_key_is_rowkey_not_rowid(self) -> None:
        """🔴 全 G 循环唯一的 `rowKey`（GC-6 裁决为最稳一族）。"""
        assert SPEC_G1402.row_identity_key == "rowKey"
        assert not SPEC_G1402.row_section_field, "固定单区不得声明 row_section_field"

    def test_footer_is_column_wise_sum_with_l20_boolean_exception(self, ws) -> None:
        """footer R20 逐列 `=SUM(x11:x19)`，🔴 `L20` 例外是布尔 `=D20=K20`。"""
        for col in ("B", "C", "D", "F", "G", "H", "I", "J", "K"):
            assert ws[f"{col}20"].value == f"=SUM({col}11:{col}19)", (
                f"{col}20 实得 {ws[f'{col}20'].value!r}"
            )
        assert ws["L20"].value == "=D20=K20", f"L20 实得 {ws['L20'].value!r}"
        assert ws["A20"].value == SPEC_G1402.footer_marker == "合计"

    def test_audit_note_rows_below_footer_are_out_of_scope(self, ws) -> None:
        assert str(ws["A21"].value or "").startswith("三、审计说明")
        assert str(ws["A23"].value or "").startswith("四、审计结论")
        assert 21 > SPEC_G1402.footer_row

    def test_iter_reads_every_row_by_rowkey(self) -> None:
        payload = [{"rowKey": "notes"}, {"rowKey": "ar"}, {"rowKey": "other"}]
        assert [rid for rid, _ in iter_store_rows(SPEC_G1402, payload)] == [
            "notes",
            "ar",
            "other",
        ]

    def test_duplicate_rowkey_fails_closed(self) -> None:
        from app.services.workpaper_sync.phase5_row_table_sheet import (
            RowTableStorePayloadError,
        )

        with pytest.raises(RowTableStorePayloadError, match="重复行身份"):
            list(iter_store_rows(SPEC_G1402, [{"rowKey": "ar"}, {"rowKey": "ar"}]))

    def test_merge_keys_rows_by_rowkey(self) -> None:
        class _FV:
            def __init__(self, key: str, row_key: str, value):
                self.stable_key = key
                self.row_key = row_key
                self.value = value
                self.is_protected = False

        key = f"{SPEC_G1402.table_key}/ar/current_unadjusted"

        class _Proj:
            def stable_keys(self):
                return [key]

            def get(self, k):
                return _FV(k, "ar", 1234) if k == key else None

        rows, applied, _visited, touched = merge_projection_into_store_rows(
            SPEC_G1402,
            projection=_Proj(),
            base_rows=[{"rowKey": "notes"}, {"rowKey": "ar"}],
        )
        assert applied == 1 and touched == {"ar"}
        hit = next(r for r in rows if r["rowKey"] == "ar")
        assert hit["currentUnadjusted"] == 1234
        # 固定行集：不新增行
        assert [r["rowKey"] for r in rows] == ["notes", "ar"]

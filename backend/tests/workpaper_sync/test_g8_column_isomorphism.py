# -*- coding: utf-8 -*-
"""G8-2 列同构判据 —— 堵住「猜测映射」+ 模板行级缺陷已修后的公式列判据（含判据 P12）。

spec: `g-cycle-single-region-detail-lanes` · Task 9b / C-8
　　　设计依据 `evidence/task8-c6-remaining-eight-template-logic.md` §3 + §12
　　　模板修正：`oo-writeback-template-parity-closure` R1（design §2 第 8 组）

═══ 四层闭环 ═══
  ① `header_text` 逐列 == **模板逐格实测**（模板是权威）
  ② `json_key` 集合/顺序 == **前端行接口**（双向相等）
  ③ 🔴 FVOCI 口径：OCI 三列必须在受管面内（**与 G9/G10 反向**）
  ④ 模板行级缺陷已修：八个公式列逐行同形 + 判据 **P12**（R/T 判 formula 时模板每行都有公式）

═══ P12 的技术根据 ═══

`excel_materialize` 写受保护格时要求 `view.has_formula`，否则抛
`ProtectedRegionWriteError`。修模板前 R13-R20 的 `R`、R12 的 `T` 整格无公式，只能判
editable；修模板后 R11-R20 每行都有 `R=F+J+L` / `T=Q+S`，于是改判 formula。
判据因此从「R/T 不得误标 formula」翻成「R/T 判 formula 且模板每个数据行都有公式」：
任一行公式丢失（模板回退），materialize 会炸，这里先红。
"""
from __future__ import annotations

import os
import re
import sys
from pathlib import Path

import openpyxl
import pytest

_REPO = Path(__file__).resolve().parents[3]
_BACKEND = _REPO / "backend"
if str(_BACKEND) not in sys.path:
    sys.path.insert(0, str(_BACKEND))
os.environ.setdefault("DB_DISABLE_SSL", "True")

from app.services.workpaper_sync.phase5_g8_02_detail import (  # noqa: E402
    FIELD_SPECS_G802,
    FOOTER_ROW_G802,
    FORMULA_COLUMNS_G802,
    FORMULA_TEMPLATES_G802,
    MANAGED_SHEET_G802,
    SPEC_G802,
    STORE_ITEM_ID_G802,
)
from app.services.workpaper_sync.phase5_row_table_sheet import (  # noqa: E402
    iter_store_rows,
    merge_projection_into_store_rows,
)

TPL = _BACKEND / "wp_templates" / "G" / "G8 其他权益工具投资.xlsx"
FE_COMPOSABLE = (
    _REPO
    / "audit-platform/frontend/src/components/workpaper/composables/useG8Detail.ts"
)

#: 前端行接口里**不受管**的字段（身份 / 显示序号）。G8 单区 ⇒ 无 `section`。
_NON_MANAGED_FE_FIELDS = {"rowId", "seq"}

_DATA_ROWS = tuple(range(SPEC_G802.first_data_row, SPEC_G802.last_data_row + 1))


@pytest.fixture(scope="module")
def ws():
    wb = openpyxl.load_workbook(TPL, data_only=False)
    try:
        yield wb[MANAGED_SHEET_G802]
    finally:
        wb.close()


def _group_start_of(ws, col_index: int) -> str:
    """该列所属的 R9 **横向**分组起始格；不属任何分组则 `""`。"""
    for rng in ws.merged_cells.ranges:
        if (
            rng.min_row == 9
            and rng.max_row == 9
            and rng.max_col > rng.min_col
            and rng.min_col <= col_index <= rng.max_col
        ):
            return rng.coord.split(":")[0]
    return ""


def _frontend_row_fields() -> list[str]:
    src = FE_COMPOSABLE.read_text(encoding="utf-8")
    m = re.search(r"export interface G8DetailRow \{(?P<body>.*?)\n\}", src, re.S)
    assert m, "找不到 G8DetailRow 接口 —— 前端结构变了，判据须改写"
    return [
        fm.group(1)
        for line in m.group("body").splitlines()
        if (fm := re.match(r"\s*(\w+)\??:", line))
    ]


# ════════════════════════════════════════════════════════════════════════════
# ① header_text 逐列 == 模板逐格实测
# ════════════════════════════════════════════════════════════════════════════
class TestHeaderTextMatchesTemplateCellByCell:
    def test_23_columns_in_excel_order_a_to_w(self) -> None:
        from openpyxl.utils import get_column_letter

        cols = [f[1] for f in FIELD_SPECS_G802]
        assert cols == [get_column_letter(i) for i in range(1, 24)], (
            f"列序必须是 A..W 连续 23 列，实得 {cols}"
        )

    def test_effective_column_count_is_23_and_uuid_col_is_x(self, ws) -> None:
        """🔴 `max_column=24` 含空列 —— 有效列 23，UUID 列据此取 X 而非 Y。"""
        from openpyxl.utils import get_column_letter

        effective = [
            get_column_letter(c)
            for c in range(1, ws.max_column + 1)
            if any(
                ws.cell(row=r, column=c).value not in (None, "")
                for r in (9, 10, *_DATA_ROWS, FOOTER_ROW_G802)
            )
        ]
        assert effective == [get_column_letter(i) for i in range(1, 24)], (
            f"有效内容列实测 {effective}"
        )
        assert ws.max_column == 24
        assert SPEC_G802.uuid_col == "X"

    @pytest.mark.parametrize(
        "field", FIELD_SPECS_G802, ids=[f[1] for f in FIELD_SPECS_G802]
    )
    def test_header_text_equals_template_cell(self, field, ws) -> None:
        """落在 R9 横向合并区内的列取叶子行 R10，跨两行合并的单列取 R9。"""
        from openpyxl.utils import column_index_from_string

        _key, col, _mode, _vt, _json_key, header_text, _group = field
        ci = column_index_from_string(col)
        leaf = ws.cell(row=10, column=ci).value
        group = ws.cell(row=9, column=ci).value
        expect = (
            str(leaf).strip()
            if (leaf and str(leaf).strip())
            else str(group or "").strip()
        )
        assert header_text == expect, (
            f"{col} 列 header_text 实得 {header_text!r}，模板逐格为 {expect!r}"
            f"（R9={group!r} / R10={leaf!r}）"
        )

    @pytest.mark.parametrize(
        "field", FIELD_SPECS_G802, ids=[f[1] for f in FIELD_SPECS_G802]
    )
    def test_group_header_cell_follows_r9_merge_geometry(self, field, ws) -> None:
        from openpyxl.utils import column_index_from_string

        _key, col, _mode, _vt, _json_key, _header, group_cell = field
        expect = _group_start_of(ws, column_index_from_string(col))
        assert group_cell == expect, (
            f"{col} 列声明 group_header_cell={group_cell!r}，"
            f"按 R9 横向合并区实测应为 {expect!r}"
        )

    def test_exactly_three_r9_horizontal_groups(self, ws) -> None:
        starts = {
            rng.coord.split(":")[0]
            for rng in ws.merged_cells.ranges
            if rng.min_row == 9 and rng.max_row == 9 and rng.max_col > rng.min_col
        }
        assert sorted(starts) == ["C9", "I9", "O9"], f"R9 横向分组实得 {sorted(starts)}"
        assert {f[6] for f in FIELD_SPECS_G802 if f[6]} == starts


# ════════════════════════════════════════════════════════════════════════════
# ② json_key 集合/顺序 == 前端行接口
# ════════════════════════════════════════════════════════════════════════════
class TestJsonKeysMatchFrontendRowInterface:
    def test_declared_json_keys_equal_frontend_managed_fields(self) -> None:
        declared = {f[4] for f in FIELD_SPECS_G802}
        frontend = set(_frontend_row_fields()) - _NON_MANAGED_FE_FIELDS
        assert declared == frontend, (
            "声明的 json_key 与前端 G8DetailRow 字段必须**双向相等**。\n"
            f"  前端有而声明缺：{sorted(frontend - declared)}\n"
            f"  声明有而前端缺：{sorted(declared - frontend)}"
        )
        assert len(declared) == 23

    def test_frontend_field_order_matches_excel_column_order(self) -> None:
        fe = [f for f in _frontend_row_fields() if f not in _NON_MANAGED_FE_FIELDS]
        assert fe == [f[4] for f in FIELD_SPECS_G802], (
            "前端字段顺序应与模板列序一致（便于逐列核对）。\n"
            f"  前端：{fe}\n  声明：{[f[4] for f in FIELD_SPECS_G802]}"
        )

    def test_row_identity_is_not_a_managed_column_and_no_section(self) -> None:
        declared = {f[4] for f in FIELD_SPECS_G802}
        for f in _NON_MANAGED_FE_FIELDS:
            assert f not in declared
        assert SPEC_G802.row_identity_key == "rowId"
        assert "section" not in _frontend_row_fields(), (
            "G8 是单区，前端行接口不应有 section 字段（那是 G9 三区的形态）"
        )
        assert not SPEC_G802.row_section_field, "单区不得声明 row_section_field"

    def test_dropped_legacy_fields_are_not_managed_and_each_has_a_home(self) -> None:
        src = FE_COMPOSABLE.read_text(encoding="utf-8")
        m = re.search(
            r"DROPPED_LEGACY_G8_FIELDS[^=]*=\s*\[(?P<body>.*?)\n\]", src, re.S
        )
        assert m, "找不到 DROPPED_LEGACY_G8_FIELDS —— 迁移台账被删了"
        dropped = re.findall(r"field:\s*'(\w+)'", m.group("body"))
        assert len(dropped) == 8, f"移除字段实得 {len(dropped)} 条：{dropped}"
        declared = {f[4] for f in FIELD_SPECS_G802}
        assert not (set(dropped) & declared)
        # 五列公允价值测试族的归属必须逐字指名 G8-4
        assert "G8-4" in m.group("body")
        # 🔴 与 G9/G10 反向：移除清单里**不得**出现 OCI 三列
        for oci in ("ociOpeningCumulative", "ociToRetainedEarnings"):
            assert oci not in dropped, (
                f"{oci} 被列入移除清单 —— G8 是 FVOCI，OCI 列是 CAS22 要求的"
            )


# ════════════════════════════════════════════════════════════════════════════
# ③ 🔴 FVOCI 口径（与 G9/G10 反向）
# ════════════════════════════════════════════════════════════════════════════
class TestFvociColumnsAreRequiredNotRemovable:
    def test_three_oci_columns_are_managed(self) -> None:
        """F 期初累计 / L 本期转留存 / R 期末累计 —— 三列都必须在受管面内。"""
        by_col = {f[1]: f for f in FIELD_SPECS_G802}
        assert by_col["F"][4] == "openingOciCumulative"
        assert by_col["L"][4] == "ociToRetainedEarnings"
        assert by_col["R"][4] == "closingOciCumulative"

    def test_oci_header_texts_are_taken_verbatim_from_template(self, ws) -> None:
        by_col = {f[1]: f for f in FIELD_SPECS_G802}
        assert by_col["F"][5] == "计入其他综合收益的累计利得或损失"
        assert by_col["R"][5] == "计入其他综合收益的累计利得或损失"
        assert by_col["L"][5] == "其他综合收益转入留存收益"
        # 逐格复核（F 与 R 同名，靠列位区分 —— 判据必须两格都核）
        assert ws["F10"].value == by_col["F"][5]
        assert ws["R10"].value == by_col["R"][5]
        assert ws["L10"].value == by_col["L"][5]

    def test_fvoci_designation_is_irrevocable_per_template_note(self, ws) -> None:
        """🔴 「该指定一经作出，不得撤销」—— 这条注释是 OCI 列不可删的依据，逐字在册。

        判据按值搜整册（注释所在 sheet 由模板决定，不写死行号）。
        """
        wb = openpyxl.load_workbook(TPL, data_only=False)
        try:
            found = False
            for sheet in wb.worksheets:
                for row in sheet.iter_rows(values_only=True):
                    for v in row:
                        if isinstance(v, str) and "不得撤销" in v and "其他综合收益" in v:
                            found = True
                            break
                    if found:
                        break
                if found:
                    break
        finally:
            wb.close()
        assert found, (
            "模板里找不到「指定为…其他综合收益…不得撤销」的注释 —— "
            "若模板真的改了口径，OCI 三列的保留依据要重新论证"
        )

    def test_dividend_income_is_a_pl_item_not_in_any_balance_group(self, ws) -> None:
        """N 本期确认的股利收入是损益项 —— 在「本期变动」组内但不进任何余额公式。"""
        by_col = {f[1]: f for f in FIELD_SPECS_G802}
        assert by_col["N"][4] == "dividendIncome"
        assert by_col["N"][6] == "I9"          # 在本期变动分组内
        assert "N" not in FORMULA_COLUMNS_G802  # 自身不是公式列
        for col, tpl in FORMULA_TEMPLATES_G802.items():
            assert "N{r}" not in tpl, f"{col} 的公式引用了 N 列：{tpl}"
        for r in _DATA_ROWS:
            m = ws[f"M{r}"].value
            assert isinstance(m, str) and "N" not in m.replace("SUM", ""), (
                f"M{r} 的本期变动合计引用了 N 列：{m!r}"
            )

    def test_movement_cost_is_net_single_column(self, ws) -> None:
        """I 本期变动/成本是**净额**单列 —— 模板没有「本期减少」叶子标题。"""
        leaves = {str(ws.cell(row=10, column=c).value or "") for c in range(9, 15)}
        assert "本期减少" not in leaves, f"I..N 叶子标题实得 {leaves}"
        declared = {f[4] for f in FIELD_SPECS_G802}
        assert "decreaseAmount" not in declared


# ════════════════════════════════════════════════════════════════════════════
# ④ 模板行级缺陷已修：八个公式列逐行同形 + 判据 P12
# ════════════════════════════════════════════════════════════════════════════
class TestTemplateRowLevelFormulasAndP12:
    def test_formula_columns_are_exactly_the_columns_with_formula_on_every_row(
        self, ws
    ) -> None:
        """判据是「模板 R11-R20 **每行都有**公式」的列集合，与声明逐值相等。"""
        from openpyxl.utils import get_column_letter

        every_row = tuple(
            get_column_letter(c)
            for c in range(1, 24)
            if all(
                isinstance(ws.cell(row=r, column=c).value, str)
                and str(ws.cell(row=r, column=c).value).startswith("=")
                for r in _DATA_ROWS
            )
        )
        assert FORMULA_COLUMNS_G802 == every_row, (
            f"formula_columns 声明 {FORMULA_COLUMNS_G802} 与「每行都有公式」的实测 "
            f"{every_row} 不一致"
        )
        assert FORMULA_COLUMNS_G802 == ("E", "H", "M", "O", "P", "Q", "R", "T")

    @pytest.mark.parametrize("col", sorted(FORMULA_TEMPLATES_G802))
    def test_uniform_formula_template_renders_to_every_data_row(
        self, col: str, ws
    ) -> None:
        """八个公式列逐行同形 ⇒ 每个数据行逐格等于声明模板渲染结果。"""
        from openpyxl.utils import column_index_from_string

        ci = column_index_from_string(col)
        for r in _DATA_ROWS:
            got = ws.cell(row=r, column=ci).value
            assert got == FORMULA_TEMPLATES_G802[col].format(r=r), (
                f"{col}{r} 模板实得 {got!r}，声明渲染为 "
                f"{FORMULA_TEMPLATES_G802[col].format(r=r)!r}"
            )

    def test_every_formula_column_is_in_the_uniform_template_table(self) -> None:
        """修模板前 M/P 逐行不同形、只能记逐行台账；修后八列全进统一模板表，台账删除。"""
        assert set(FORMULA_COLUMNS_G802) == set(FORMULA_TEMPLATES_G802)
        from app.services.workpaper_sync import phase5_g8_02_detail as detail

        for stale in (
            "TEMPLATE_ROW_FORMULAS_G802",
            "TEMPLATE_ROW_DEFECTS_G802",
            "FRONTEND_DERIVED_COLUMNS_G802",
        ):
            assert not hasattr(detail, stale), (
                f"{stale} 是修模板前的缺陷台账，模板已逐行同形，不应再导出"
            )

    def test_movement_total_excludes_oci_transfer_and_closes_the_balance(self, ws) -> None:
        """🔴 M = SUM(I:K)，不含 L：E + M 必须等于 Q（L 是权益内部结转，只动 R 列 OCI 累计）。

        修模板前只有 R11 是 `SUM(I11:L11)`，其余 9 行本就是 `SUM(I:K)`。逐行核两条：
        M 的区间止于 K；按模板公式代数展开 `E + M == Q`（C+D + I+J+K == C+I + D+J+K）。
        """
        import re
        from collections import Counter

        def leaves(col: str) -> Counter:
            """把公式列按声明模板递归展开到可编辑列（模板只有加法与同行 SUM 区间）。"""
            tpl = FORMULA_TEMPLATES_G802.get(col)
            if tpl is None:
                return Counter({col: 1})
            body = tpl.format(r=1)[1:]
            out: Counter = Counter()
            for a, b in re.findall(r"SUM\(([A-Z])1:([A-Z])1\)", body):
                for code in range(ord(a), ord(b) + 1):
                    out += leaves(chr(code))
            for c in re.findall(r"([A-Z])1", re.sub(r"SUM\([A-Z]1:[A-Z]1\)", "", body)):
                out += leaves(c)
            return out

        for r in _DATA_ROWS:
            assert ws[f"M{r}"].value == f"=SUM(I{r}:K{r})", f"M{r} 实得 {ws[f'M{r}'].value!r}"
        assert leaves("E") + leaves("M") == leaves("Q"), (
            f"E+M 展开 {sorted((leaves('E') + leaves('M')).elements())} ≠ "
            f"Q 展开 {sorted(leaves('Q').elements())}"
        )
        assert "L" not in leaves("M") and "L" in leaves("R")

    def test_p12_r_and_t_are_formula_and_backed_on_every_data_row(self, ws) -> None:
        """🔴 **判据 P12**：R/T 判 formula 的前提是模板每个数据行那格都有公式。

        技术根据：`excel_materialize` 写受保护格时要求 `view.has_formula`，
        缺一行公式就会抛 `ProtectedRegionWriteError`。
        """
        by_col = {f[1]: f for f in FIELD_SPECS_G802}
        assert by_col["R"][2] == "formula"
        assert by_col["T"][2] == "formula"
        for r in _DATA_ROWS:
            assert ws[f"R{r}"].value == f"=F{r}+J{r}+L{r}", f"R{r} 实得 {ws[f'R{r}'].value!r}"
            assert ws[f"T{r}"].value == f"=Q{r}+S{r}", f"T{r} 实得 {ws[f'T{r}'].value!r}"

    def test_judgement_is_discriminative_on_a_reverted_template(self, ws) -> None:
        """变异自检：把模板回退到修复前的形态（R13 的 R、R12 的 T 清空），「每行都有公式」的判据必须丢列。"""
        from openpyxl import Workbook
        from openpyxl.utils import get_column_letter

        wb = Workbook()
        mut = wb.active
        for r in _DATA_ROWS:
            for c in range(1, 24):
                mut.cell(row=r, column=c, value=ws.cell(row=r, column=c).value)
        mut["R13"] = None
        mut["T12"] = None
        every_row = tuple(
            get_column_letter(c)
            for c in range(1, 24)
            if all(
                isinstance(mut.cell(row=r, column=c).value, str)
                and str(mut.cell(row=r, column=c).value).startswith("=")
                for r in _DATA_ROWS
            )
        )
        assert every_row == ("E", "H", "M", "O", "P", "Q"), every_row
        assert every_row != FORMULA_COLUMNS_G802

    def test_every_formula_mode_field_is_covered_by_formula_mask(self) -> None:
        """契约层 CS-13 的本地复核（提前于 parse_contract 打红，定位更快）。"""
        from app.services.workpaper_sync.contracts import column_in_ranges

        mask = SPEC_G802.formula_mask
        for key, col, mode, _vt, _json, _hdr, _grp in FIELD_SPECS_G802:
            if mode != "formula":
                continue
            assert column_in_ranges(col, mask), (
                f"{key}（列 {col}）声明 mode=formula 但不在 formula_mask {mask} 内 —— "
                "会被 contracts CS-13 拒"
            )


# ════════════════════════════════════════════════════════════════════════════
# ⑤ 单区几何与 store 读写
# ════════════════════════════════════════════════════════════════════════════
class TestSingleRegionGeometryAndStore:
    def test_single_managed_region_r11_to_r20_with_footer_outside(self) -> None:
        assert (SPEC_G802.first_data_row, SPEC_G802.last_data_row) == (11, 20)
        assert SPEC_G802.footer_row == FOOTER_ROW_G802 == 21
        assert SPEC_G802.footer_row > SPEC_G802.last_data_row
        assert SPEC_G802.store_item_id == STORE_ITEM_ID_G802 == "G8-detail-rows"

    def test_footer_is_column_wise_sum(self, ws) -> None:
        """footer R21 逐列 `=SUM(x11:x20)`（G8 无 G10 那样的 O21 例外）。

        修模板前 K / L / R 三列合计格是空的（`oo-writeback-template-parity-closure` R1 补上），
        现在 C..T 十八列都必须是同形 SUM。
        """
        for col in "CDEFGHIJKLMNOPQRST":
            assert ws[f"{col}21"].value == f"=SUM({col}11:{col}20)", (
                f"{col}21 实得 {ws[f'{col}21'].value!r}"
            )
        assert ws["A21"].value == SPEC_G802.footer_marker == "合计"

    def test_audit_note_row_below_footer_is_out_of_managed_scope(self, ws) -> None:
        """🔴 R22「三、审计说明：」有个 `N22==A22` 的模板怪癖 —— 必须落在受管区外。"""
        assert str(ws["A22"].value or "").startswith("三、审计说明")
        assert ws["N22"].value == "=A22"
        assert 22 > SPEC_G802.footer_row

    def test_ghost_row_anchor_uses_default_column_a(self) -> None:
        """G8 的 A 列就是业务名称本身（模板 R11 无预填值）⇒ 用默认锚点 0，不像 G9/G10 改指 B。"""
        assert SPEC_G802.ghost_row_anchor_index == 0
        anchor = FIELD_SPECS_G802[0]
        assert (anchor[1], anchor[4]) == ("A", "investeeName")
        from openpyxl import load_workbook

        wb = load_workbook(TPL, data_only=False)
        try:
            assert wb[MANAGED_SHEET_G802]["A11"].value in (None, ""), (
                "模板 A11 有预填值 ⇒ 幽灵行锚点要像 G9/G10 那样改指别的列"
            )
        finally:
            wb.close()

    def test_iter_reads_every_row_without_section_filtering(self) -> None:
        payload = [
            {"rowId": "a1", "investeeName": "甲公司"},
            {"rowId": "a2", "investeeName": "乙公司"},
        ]
        assert [rid for rid, _ in iter_store_rows(SPEC_G802, payload)] == ["a1", "a2"]

    def test_duplicate_identity_fails_closed(self) -> None:
        from app.services.workpaper_sync.phase5_row_table_sheet import (
            RowTableStorePayloadError,
        )

        with pytest.raises(RowTableStorePayloadError, match="重复行身份"):
            list(iter_store_rows(SPEC_G802, [{"rowId": "x"}, {"rowId": "x"}]))

    def test_merge_does_not_stamp_section_on_new_rows(self) -> None:
        class _FV:
            def __init__(self, key: str, row_key: str, value):
                self.stable_key = key
                self.row_key = row_key
                self.value = value
                self.is_protected = False

        new_key = f"{SPEC_G802.table_key}/newrow/investee_name"

        class _Proj:
            def stable_keys(self):
                return [new_key]

            def get(self, key):
                return _FV(key, "newrow", "OO侧新增") if key == new_key else None

        rows, applied, _visited, touched = merge_projection_into_store_rows(
            SPEC_G802, projection=_Proj(), base_rows=[{"rowId": "a1", "investeeName": "甲"}]
        )
        assert applied == 1 and touched == {"newrow"}
        added = next(r for r in rows if r["rowId"] == "newrow")
        assert added["investeeName"] == "OO侧新增"
        assert "section" not in added

"""P2 / P3 真链判据：D4-1 store → **真物化** → **真提取** → merge → HTML 读回。

spec: d4-html-to-oo-store-contract-alignment · Task 11（真链 materialize 段）· Task 17
Requirements 1.2 / 1.3 / 2.1 · Property 2 / Property 3

═══ 为什么必须有这一条（它补的是一个真缺口，不是重复覆盖）═══

Task 11 原先的 P4 判据（`test_d4_1_oo_to_html_realchain.py`）用「在 projection 里把
`FieldValue.value` 改掉」代替 OO 真产物 —— 于是 **materialize / extract 那一段从未被判据覆盖**。
而需求 1.2 原文是「这条挡的是 *projection 对了但 binding 没写*」：只在 projection 上改值，
恰恰绕过了 binding 这一环，挡不住它要挡的东西。

本文件走真 xlsx 字节：

  1. instrument 真模板（`D4.instrumentation_specs()`）得 substrate，行 UUID 由模板占位提供
     （`GTROW-D41MAIN-0008…0011` / `GTROW-D41OTHER-0014…0017`，与缺陷现场同一批占位）；
  2. `build_store_projection_d41(新形态行对象)` —— 金额**随行落库**（Task 8 的形态）；
  3. `adapter.materialize()` → 真 xlsx 字节。D4-1 是**同 sheet 双区**，故 primary binding
     取主营、其他段走 `sibling_bindings`（与生产 `_sibling_identity_bindings` 同形态）；
  4. `adapter.extract()` → 反读 projection；
  5. **P2**：3/4 两步逐字段等值；
  6. **P3**：48 格 `FORMULA_MASK` 在 projection 与 extract 产物两侧都恒不产普通值键；
  7. **T11 补段**：openpyxl 直接改 staged 产物的 `B8`（主营首行本期未审数，模拟审计师在 OO
     里改数）→ extract → `merge_projection_into_d41_rows` → 断言行对象顶层 store_key 等于
     改后值（这就是前端 `readRowFieldWithFallback` 会读到的值）。

🔴 **金额比较必须 `float()` 归一，不得裸 `==`**：实测 openpyxl 往返把 `100.0` 读回成 `100`、
`11000.0` 读回成 `11000`。平台已在 projection digest 上踩过同款表示分叉（`0` vs `0.0` 使复用
恒不命中），这里用同一条口径，不是放宽判据。
"""
from __future__ import annotations

import hashlib
import os
import sys
import uuid
from pathlib import Path
from typing import Any, Mapping

import pytest

_REPO = Path(__file__).resolve().parents[3]
_BACKEND = _REPO / "backend"
if str(_BACKEND) not in sys.path:  # pragma: no cover
    sys.path.insert(0, str(_BACKEND))
os.environ.setdefault("DB_DISABLE_SSL", "True")

from app.services.excel_structure_fingerprint import identity_inventory  # noqa: E402
from app.services.workpaper_sync import excel_extract as X  # noqa: E402
from app.services.workpaper_sync import excel_instrumentation as EI  # noqa: E402
from app.services.workpaper_sync import phase5_d4_adjudication_sheet as A  # noqa: E402
from app.services.workpaper_sync import phase5_d4_revenue_detail as D4  # noqa: E402
from app.services.workpaper_sync.adapters import excel as AX  # noqa: E402
from app.services.workpaper_sync.contracts import parse_contract  # noqa: E402
from app.services.workpaper_sync.excel_entry_gate import (  # noqa: E402
    AdapterBuild,
    FrozenEntryDefinitions,
    parse_identity_inventory,
)
from app.services.workpaper_sync.models import (  # noqa: E402
    AuthorityModel,
    BundleSlot,
    BundleSlotSpec,
    DefinitionState,
)
from app.services.workpaper_sync.resolution import DefinitionBundleSnapshot  # noqa: E402

#: 6 个受管金额的 store_key（行对象顶层键）。
_AMOUNT_KEYS = (
    "currentUnadjusted",
    "currentAje",
    "currentRje",
    "priorUnadjusted",
    "priorAje",
    "priorRje",
)
#: store_key → column_key（stable_key 用 column_key）。
_STORE_TO_COL = {v: k for k, v in A._COLUMN_KEY_TO_STORE_KEY.items()}


# ═══════════════════════════════════════════════════════════════════════════
# Fixtures：真模板 instrument + 双区 binding 的 adapter（照 test_d4_14 范式）
# ═══════════════════════════════════════════════════════════════════════════


def _d(label: str) -> str:
    return hashlib.sha256(label.encode("utf-8")).hexdigest()


def _make_bundle(contract: Any) -> DefinitionBundleSnapshot:
    def slot(kind: BundleSlot, digest: str) -> BundleSlotSpec:
        return BundleSlotSpec(kind, "definition", f"definition:{uuid.uuid4()}", digest)

    return DefinitionBundleSnapshot(
        bundle_id=uuid.uuid4(),
        bundle_sha256=_d("d41-realchain-bundle"),
        schema_version="definition-bundle:v1",
        state=DefinitionState.approved,
        authority_model=AuthorityModel.projection_contract,
        authority_model_definition_id=uuid.uuid4(),
        authority_model_definition_sha256=_d("d41-realchain-authority"),
        slots={
            BundleSlot.template: slot(BundleSlot.template, contract.template_definition_sha256),
            BundleSlot.instrumentation: slot(
                BundleSlot.instrumentation, contract.instrumentation_definition_sha256
            ),
            BundleSlot.contract: slot(BundleSlot.contract, contract.canonical_sha256),
        },
    )


def _make_definitions(contract: Any, inventory_raw: Mapping[str, Any]) -> FrozenEntryDefinitions:
    return FrozenEntryDefinitions(
        entry_id=D4.ENTRY_ID,
        bundle=_make_bundle(contract),
        contract=contract,
        adapter_build=AdapterBuild(
            adapter_id=D4.ADAPTER_ID,
            adapter_build_digest=_d("d41-realchain-adapter-build"),
            document_type="xlsx",
            contract_version=contract.semantic_version,
        ),
        identity_inventory=parse_identity_inventory(inventory_raw),
        business_sheets=(),
        dynamic_column_keys={},
        structure_inventory_size=0,
    )


BINDING_MAIN = X.ExcelIdentityBinding(
    table_name=A.TABLE_NAME_MAIN,
    uuid_column=A.UUID_COL_MAIN,
    table_key=A.ROWS_TABLE_KEY_MAIN,
)
BINDING_OTHER = X.ExcelIdentityBinding(
    table_name=A.TABLE_NAME_OTHER,
    uuid_column=A.UUID_COL_OTHER,
    table_key=A.ROWS_TABLE_KEY_OTHER,
)


@pytest.fixture(scope="module")
def contract() -> Any:
    return parse_contract(D4.build_contract_payload(), adapter_id=D4.ADAPTER_ID)


@pytest.fixture(scope="module")
def instrumented() -> EI.InstrumentedWorkbook:
    source = D4.read_authoritative_template()
    gate = D4.excel_carrier_gate()
    return EI.instrument_workbook_bytes_multi(source, D4.instrumentation_specs(), gate=gate)


@pytest.fixture(scope="module")
def seed_uuids(instrumented: EI.InstrumentedWorkbook) -> tuple[str, str]:
    """两区首数据行的模板占位 UUID（与缺陷现场 overlay 退回的那批占位同源）。"""
    inv_m = identity_inventory(
        instrumented.instrumented_bytes,
        expected_table=A.TABLE_NAME_MAIN,
        uuid_column_letter=A.UUID_COL_MAIN,
    )
    inv_o = identity_inventory(
        instrumented.instrumented_bytes,
        expected_table=A.TABLE_NAME_OTHER,
        uuid_column_letter=A.UUID_COL_OTHER,
    )
    um = inv_m["hidden_uuid_column"]["row_uuids"]
    uo = inv_o["hidden_uuid_column"]["row_uuids"]
    rid_m = um.get(str(A.FIRST_DATA_ROW_MAIN)) or um.get(A.FIRST_DATA_ROW_MAIN)
    rid_o = uo.get(str(A.FIRST_DATA_ROW_OTHER)) or uo.get(A.FIRST_DATA_ROW_OTHER)
    assert rid_m and rid_o, f"D4-1 两区 seed UUID 缺失: main={um} other={uo}"
    return str(rid_m), str(rid_o)


@pytest.fixture(scope="module")
def definitions(contract: Any, instrumented: EI.InstrumentedWorkbook) -> FrozenEntryDefinitions:
    inv = identity_inventory(
        instrumented.instrumented_bytes,
        expected_table=A.TABLE_NAME_MAIN,
        uuid_column_letter=A.UUID_COL_MAIN,
    )
    return _make_definitions(contract, inv)


@pytest.fixture
def base_path(instrumented: EI.InstrumentedWorkbook, tmp_path: Path) -> Path:
    path = tmp_path / "d41-base.xlsx"
    path.write_bytes(instrumented.instrumented_bytes)
    return path


@pytest.fixture
def adapter(definitions: FrozenEntryDefinitions) -> Any:
    """双区 adapter：主营为 primary、其他段走 sibling_bindings（同生产路径形态）。"""
    return AX.build_excel_adapter(
        definitions=definitions,
        binding=BINDING_MAIN,
        direction="html_to_oo",
        sibling_bindings=(BINDING_OTHER,),
    )


def _html_rows(rid_main: str, rid_other: str) -> list[dict[str, Any]]:
    """Task 8 之后的前端落库形态：金额**随行**平铺进行对象顶层（number，不是字符串）。"""
    return [
        {
            "rowId": rid_main,
            "label": "主营-批发分销",
            "source": "tb",
            "accountCode": "6001",
            A.SECTION_KEY_FIELD: A.SECTION_KEY_MAIN,
            "currentUnadjusted": 12345.67,
            "currentAje": 100.0,
            "currentRje": 0,
            "priorUnadjusted": 11000.0,
            "priorAje": 0,
            "priorRje": 0,
            # HTML-only 键（Property 5 面）：派生快照必须穿过全链不被冲掉。
            "derivedSnapshot": {"currentUnadjusted": 12345.67, "priorUnadjusted": 11000.0},
        },
        {
            "rowId": rid_other,
            "label": "其他-废料销售",
            "source": "manual",
            "accountCode": "6051",
            A.SECTION_KEY_FIELD: A.SECTION_KEY_OTHER,
            "currentUnadjusted": 777.77,
            "currentAje": 0,
            "currentRje": 0,
            "priorUnadjusted": 0,
            "priorAje": 0,
            "priorRje": 0,
        },
    ]


def _num(value: Any) -> float | None:
    """金额归一：openpyxl 往返会把 100.0 读回成 100，裸 `==` 会假红（既有表示分叉）。"""
    if value is None or value == "":
        return None
    return float(value)


def _stage(adapter: Any, contract: Any, base: Path, proj: Any, tmp: Path, name: str) -> Path:
    staged = tmp / AX.STAGING_NAMESPACE / name
    staged.parent.mkdir(parents=True, exist_ok=True)
    adapter.materialize(substrate=base, projection=proj, output=staged, contract=contract)
    assert staged.is_file() and staged.read_bytes()[:2] == b"PK", "materialize 未产出 xlsx 字节"
    return staged


# ═══════════════════════════════════════════════════════════════════════════
# Property 2：projection ≡ materialize 后 extract 的反读结果（逐字段）
# **Validates: Requirements 1.2**
# ═══════════════════════════════════════════════════════════════════════════


def test_p2_projection_equals_extract_field_by_field(
    contract: Any, adapter: Any, base_path: Path, seed_uuids: tuple[str, str], tmp_path: Path
) -> None:
    """P2：新存储形态下，物化产物反读出的两区字段集合与值，与 projection 逐项相等。

    这条挡「projection 对了但 binding 没写」—— 在 projection 上改值的判据对此天生是盲的。
    """
    rid_m, rid_o = seed_uuids
    rows = _html_rows(rid_m, rid_o)
    proj = A.build_store_projection_d41(rows, contract=contract)
    assert proj.row_keys[A.ROWS_TABLE_KEY_MAIN] == (rid_m,)
    assert proj.row_keys[A.ROWS_TABLE_KEY_OTHER] == (rid_o,)

    staged = _stage(adapter, contract, base_path, proj, tmp_path, "d41-p2.xlsx")
    extracted = adapter.extract(artifact=staged, contract=contract)

    # 1) projection 的每一个键都必须在 extract 产物里，且值相等（金额 float 归一、文本逐字）。
    missing: list[str] = []
    mismatch: list[str] = []
    for key, fv in proj.values.items():
        got = extracted.get(key)
        if got is None:
            missing.append(str(key))
            continue
        want, have = fv.value, got.value
        if isinstance(want, (int, float)) or isinstance(have, (int, float)):
            if _num(want) != _num(have):
                mismatch.append(f"{key}: projection={want!r} extract={have!r}")
        elif str(want) != str(have):
            mismatch.append(f"{key}: projection={want!r} extract={have!r}")
    assert not missing, f"materialize 写了但 extract 读不回（binding 漏写）: {missing}"
    assert not mismatch, f"反读值与 projection 不等: {mismatch}"

    # 2) 反向：extract 不得凭空多出本次投影之外的受管两区键（空占位行不产值键）。
    two_region_extract = {
        str(k)
        for k in extracted.values
        if str(k).startswith((A.ROWS_TABLE_KEY_MAIN + "/", A.ROWS_TABLE_KEY_OTHER + "/"))
    }
    assert two_region_extract == {str(k) for k in proj.values}, (
        f"两区键集合不等：extract 多出 {sorted(two_region_extract - {str(k) for k in proj.values})}"
    )
    # 3) 14 格 = 2 行 × (label + 6 金额)，钉住规模不被悄悄缩小。
    assert len(two_region_extract) == 14, f"两区受管格应为 14，实得 {len(two_region_extract)}"


def test_p2_html_only_keys_are_not_materialized_as_managed(
    contract: Any, adapter: Any, base_path: Path, seed_uuids: tuple[str, str], tmp_path: Path
) -> None:
    """`derivedSnapshot` 等 HTML-only 键 SHALL NOT 变成受管格（它只活在 store 行对象里）。

    若哪天有人把它接进 `MANAGED_FIELD_SPECS`，它就会被写进 Excel 某列、并被 OO 侧编辑 ——
    覆盖状态机的第三个量随即失去「只由派生写入」的语义，S2/S4 判定全盘失真。
    """
    rid_m, rid_o = seed_uuids
    proj = A.build_store_projection_d41(_html_rows(rid_m, rid_o), contract=contract)
    staged = _stage(adapter, contract, base_path, proj, tmp_path, "d41-htmlonly.xlsx")
    extracted = adapter.extract(artifact=staged, contract=contract)
    for key in list(proj.values) + list(extracted.values):
        sk = str(key)
        assert "derivedSnapshot" not in sk and "derived_snapshot" not in sk, (
            f"HTML-only 键泄漏成受管格: {sk}"
        )
        assert not sk.endswith(("/source", "/accountCode", "/sectionKey")), (
            f"结构键泄漏成受管格: {sk}"
        )


# ═══════════════════════════════════════════════════════════════════════════
# Property 3：FORMULA_MASK 48 格恒不产普通值键（新存储形态下不放宽）
# **Validates: Requirements 1.3**
# ═══════════════════════════════════════════════════════════════════════════


def test_p3_formula_mask_scale_is_48_and_covers_audited_and_footer() -> None:
    """48 格规模与覆盖面不因「金额随行落库」而被放宽（本 spec 自带判据，不只依赖归档 spec）。"""
    mask = set(A.FORMULA_MASK)
    assert len(A.FORMULA_MASK) == 48, f"FORMULA_MASK 应为 48 格，实得 {len(A.FORMULA_MASK)}"
    # E/I 审定数（=SUM）两区数据行首末。
    for row in (A.FIRST_DATA_ROW_MAIN, A.LAST_DATA_ROW_MAIN,
                A.FIRST_DATA_ROW_OTHER, A.LAST_DATA_ROW_OTHER):
        for col in ("E", "I"):
            assert f"{col}{row}" in mask, f"审定数格 {col}{row} 掉出 mask"
    # 小计 / 合计 / 差异行的 B–I。
    for row in (A.SUBTOTAL_ROW_MAIN, A.SUBTOTAL_ROW_OTHER, 19, 21):
        for col in ("B", "C", "D", "E", "F", "G", "H", "I"):
            assert f"{col}{row}" in mask, f"公式行格 {col}{row} 掉出 mask"


def test_p3_masked_cells_never_produce_plain_value_keys(
    contract: Any, adapter: Any, base_path: Path, seed_uuids: tuple[str, str], tmp_path: Path
) -> None:
    """48 格在 projection 与 extract 产物**两侧**都不产普通值键。

    受管列只有 A(label) + B/C/D + F/G/H 共 7 个 column_key；E/I 与小计/合计/差异行由 Excel
    内部公式产生。若 `current_audited` 之类派生列出现在任一侧，说明公式格被普通值投影覆盖了。
    """
    rid_m, rid_o = seed_uuids
    proj = A.build_store_projection_d41(_html_rows(rid_m, rid_o), contract=contract)
    staged = _stage(adapter, contract, base_path, proj, tmp_path, "d41-mask.xlsx")
    extracted = adapter.extract(artifact=staged, contract=contract)

    managed_cols = set(_STORE_TO_COL.values())
    for side, projection in (("projection", proj), ("extract", extracted)):
        for key in projection.values:
            sk = str(key)
            if not sk.startswith((A.ROWS_TABLE_KEY_MAIN + "/", A.ROWS_TABLE_KEY_OTHER + "/")):
                continue
            col_key = sk.rsplit("/", 1)[-1]
            assert col_key in managed_cols, (
                f"[{side}] 非受管 column_key {col_key!r} 产了普通值键 {sk} —— "
                "审定数/小计等公式格被值投影覆盖（需求 1.3 破了）"
            )
            assert "audited" not in col_key, f"[{side}] 审定数派生列泄漏: {sk}"

    # 物化后 E/I 审定数格仍是公式（不是被写死的值）。
    from openpyxl import load_workbook

    wb = load_workbook(staged, data_only=False)
    ws = wb[A.MANAGED_SHEET_D41]
    for cell in (f"E{A.FIRST_DATA_ROW_MAIN}", f"I{A.FIRST_DATA_ROW_MAIN}",
                 f"E{A.FIRST_DATA_ROW_OTHER}", f"I{A.FIRST_DATA_ROW_OTHER}"):
        raw = ws[cell].value
        assert isinstance(raw, str) and raw.startswith("="), (
            f"{cell} 物化后不再是公式（实得 {raw!r}）—— 普通值投影盖掉了 =SUM"
        )
    wb.close()


# ═══════════════════════════════════════════════════════════════════════════
# Task 11 补段：真产物被改 → extract → merge → HTML 读回等值（P4 的物化段）
# **Validates: Requirements 2.1, 2.2**
# ═══════════════════════════════════════════════════════════════════════════

#: 主营首行「本期未审数」在 Excel 里的落点（column_key=current_unadjusted 的列标 B）。
_CELL_MAIN_CURRENT_UNADJUSTED = f"B{A.FIRST_DATA_ROW_MAIN}"
#: 审计师在 OO 里把它改成这个数。
_OO_EDITED_AMOUNT = 98765.43


def test_oo_edit_on_real_artifact_flows_back_to_html(
    contract: Any, adapter: Any, base_path: Path, seed_uuids: tuple[str, str], tmp_path: Path
) -> None:
    """全链：store → 真物化 → **改真 xlsx 格** → 真提取 → merge → 行对象顶层等于改后值。

    这是缺陷 A2 的完整反向锁。此前只在 projection 上改 `FieldValue` —— 那样连
    「Excel 那一格到底被写在哪里、改了之后 extract 能不能认出来」都没被验证过。
    """
    from openpyxl import load_workbook

    rid_m, rid_o = seed_uuids
    rows = _html_rows(rid_m, rid_o)
    proj = A.build_store_projection_d41(rows, contract=contract)
    staged = _stage(adapter, contract, base_path, proj, tmp_path, "d41-ooedit.xlsx")

    # 物化产物里该格应当先等于 HTML 侧的值（否则后面的"改"没有起点）。
    wb = load_workbook(staged, data_only=False)
    ws = wb[A.MANAGED_SHEET_D41]
    assert _num(ws[_CELL_MAIN_CURRENT_UNADJUSTED].value) == 12345.67, (
        f"{_CELL_MAIN_CURRENT_UNADJUSTED} 物化后应为 HTML 侧 12345.67，"
        f"实得 {ws[_CELL_MAIN_CURRENT_UNADJUSTED].value!r}"
    )
    # ── 模拟审计师在 OO 里改这一格 ──
    ws[_CELL_MAIN_CURRENT_UNADJUSTED] = _OO_EDITED_AMOUNT
    wb.save(staged)
    wb.close()

    extracted = adapter.extract(artifact=staged, contract=contract)
    key = f"{A.ROWS_TABLE_KEY_MAIN}/{rid_m}/{_STORE_TO_COL['currentUnadjusted']}"
    assert _num(extracted.get(key).value) == _OO_EDITED_AMOUNT, (
        f"extract 没认出 OO 的改动：{extracted.get(key).value!r}"
    )

    merged, applied, _visited, _touched = A.merge_projection_into_d41_rows(
        projection=extracted, base_rows=rows
    )
    assert applied > 0, "merge 未回写任何字段"
    row_main = next(r for r in merged if r["rowId"] == rid_m)
    # HTML 读侧等价：前端 readRowFieldWithFallback 行对象顶层优先 ⇒ 读到的就是这个值。
    assert _num(row_main["currentUnadjusted"]) == _OO_EDITED_AMOUNT, (
        f"HTML 读回不等于 OO 改后值：{row_main['currentUnadjusted']!r}"
    )
    # 未被改的格不得被连带改动（逐格语义，不是整行覆盖）。
    assert _num(row_main["priorUnadjusted"]) == 11000.0
    assert _num(row_main["currentAje"]) == 100.0
    # 其他段那一行完全不受影响。
    row_other = next(r for r in merged if r["rowId"] == rid_o)
    assert _num(row_other["currentUnadjusted"]) == 777.77

    # Property 5：HTML-only 字段穿过真链不丢（含覆盖状态机依赖的 derivedSnapshot）。
    assert row_main["source"] == "tb"
    assert row_main["accountCode"] == "6001"
    assert row_main["derivedSnapshot"] == {
        "currentUnadjusted": 12345.67,
        "priorUnadjusted": 11000.0,
    }, f"derivedSnapshot 被真链冲掉了：{row_main.get('derivedSnapshot')!r}"

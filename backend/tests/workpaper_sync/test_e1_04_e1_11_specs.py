# -*- coding: utf-8 -*-
"""E1-4（最小行表）+ E1-11（唯一 static_region）声明判据（E1 Tasks 13 / 14）。

spec: e1-sync-coverage-and-first-canary · Requirements 2.1 / 2.2 / 2.4 / 2.6

═══ 判据面 ═══

* **E1-4：第二张复用框架层零改动**（Task 13 的核心断言）——
  本模块只有一个 `RowTableSheetSpec` 常量，零算法、零框架层改动。
* **E1-11：E1-P3 转绿**（Task 14）——
  ①它是 `static_region`：`table_name`/`uuid_col` 必空、`defined_name` 必填、无行维度
  ②它**绕开整条位移链**：不产生 row_shift、不经 footer 两门、不注 UUID 列、不建 Excel Table
  ③反证：E1-9/E1-10 **不是** static_region（有 -rows 键 ⇒ 动态行表）
* **声明 ≡ 模板实测几何**（禁推演）。
* **OCR 第二写入方已登记**（E1-P16 的可执行落点）。
"""
from __future__ import annotations

import glob
import os
import re
import sys
from pathlib import Path

import pytest

_BACKEND = Path(__file__).resolve().parents[2]
if str(_BACKEND) not in sys.path:  # pragma: no cover
    sys.path.insert(0, str(_BACKEND))
os.environ.setdefault("DB_DISABLE_SSL", "True")

from app.services.workpaper_sync import phase5_e1_04_digital as E104
from app.services.workpaper_sync import phase5_e1_11_commitment as E111
from app.services.workpaper_sync import phase5_row_table_sheet as ENGINE
from app.services.workpaper_sync.excel_extract import BindingKind


@pytest.fixture(scope="module")
def e1_first_volume():
    import openpyxl

    cands = [p for p in glob.glob(str(_BACKEND / "wp_templates" / "E" / "*.xlsx"))
             if "E1-1至E1-11" in p]
    assert cands
    return openpyxl.load_workbook(cands[0], data_only=False)


# ═══════════════════════════════════════════════════════════════════════════
# E1-4：最小行表 + 零框架层改动
# ═══════════════════════════════════════════════════════════════════════════


def test_e104_geometry_matches_template(e1_first_volume) -> None:
    ws = e1_first_volume[E104.MANAGED_SHEET_E104]
    assert ws.cell(row=E104.FOOTER_ROW_E104, column=1).value == E104.FOOTER_MARKER_E104
    # 数据区 7 行，A 列预填序号 1-7
    for row, seq in E104.PREFILLED_SEQ_ROWS_E104:
        assert ws.cell(row=row, column=1).value == seq, f"R{row} 序号应为 {seq}"
    assert E104.LAST_DATA_ROW_E104 - E104.FIRST_DATA_ROW_E104 + 1 == 7


def test_e104_formula_templates_match_template(e1_first_volume) -> None:
    """H/I/K/L 四列公式模板逐字相等；I/L 是乘法（原币 × 汇率）。"""
    from openpyxl.utils import column_index_from_string

    ws = e1_first_volume[E104.MANAGED_SHEET_E104]
    for col, tpl in E104.FORMULA_TEMPLATES_E104.items():
        found = False
        for r in range(E104.FIRST_DATA_ROW_E104, E104.LAST_DATA_ROW_E104 + 1):
            v = ws.cell(row=r, column=column_index_from_string(col)).value
            if isinstance(v, str) and v.startswith("="):
                assert re.sub(rf"(?<=[A-Z]){r}\b", "{r}", v) == tpl, (
                    f"{col}{r} 实测 {v!r}，声明 {tpl!r}"
                )
                found = True
        assert found, f"{col} 列数据区无公式但声明为 formula"
    assert "*" in E104.FORMULA_TEMPLATES_E104["I"]
    assert "*" in E104.FORMULA_TEMPLATES_E104["L"]


def test_e104_module_is_declaration_only() -> None:
    """🔴 Task 13 的核心断言：第二张复用框架层**零改动** —— 声明模块里无任何算法。

    判据：模块源码里不得有 `def ` 函数定义（除 dataclass 自动生成的），
    只允许常量赋值与 import。若接第二张需要写函数，说明框架层抽象不足。
    """
    src = Path(E104.__file__).read_text(encoding="utf-8")
    # 去掉 docstring 与注释后不应有 def / class
    body = re.sub(r'"""[\s\S]*?"""', "", src)
    body = "\n".join(
        line for line in body.splitlines() if not line.strip().startswith("#")
    )
    assert "def " not in body, "声明模块不应含函数定义（算法归框架层）"
    assert "class " not in body, "声明模块不应含类定义（数据类归框架层）"


def test_e104_is_dynamic_row_table() -> None:
    assert E104.SPEC_E104.binding_kind is BindingKind.excel_table
    assert E104.SPEC_E104.uuid_col == "R"
    assert E104.SPEC_E104.row_identity_key == "id"  # E1 侧统一用 id
    assert E104.SPEC_E104.formula_mask == ("H10:H16", "I10:I16", "K10:K16", "L10:L16")


# ═══════════════════════════════════════════════════════════════════════════
# E1-11：唯一 static_region + 绕开位移链（E1-P3 转绿）
# ═══════════════════════════════════════════════════════════════════════════


def test_e111_is_static_region_with_binding_kind_rules() -> None:
    """static_region 的分派铁律：table_name/uuid_col 必空、defined_name 必填、无行维度。"""
    spec = E111.SPEC_E111
    assert spec.binding_kind is BindingKind.static_region
    assert spec.table_name == "", "static_region 的 table_name 必空"
    assert spec.uuid_col == "", "static_region 的 uuid_col 必空"
    assert spec.defined_name == E111.DEFINED_NAME_E111, "static_region 的 defined_name 必填"
    assert spec.row_identity_key == "", "static_region 无行维度"


def test_e111_bypasses_the_row_shift_chain() -> None:
    """🔴 E1-P3 转绿的核心：static_region **绕开整条位移链**。

    三条可断言的后果（声明层能验的部分）：
      ①无 Excel Table ⇒ 不会有 sibling table ref 需要位移
      ②无 UUID 列 ⇒ 不会 mint UUID、不会有隐藏列注入
      ③无 formula_columns ⇒ 不产生列向 mask（承诺函正文无逐行公式）
    运行时的 `_plan_static_writes` 按绝对坐标直写（不经 row_shift / footer 两门），
    属真栈段，卡 adapter 注册（upstream_gap）。
    """
    spec = E111.SPEC_E111
    assert not spec.table_name
    assert not spec.uuid_col
    assert spec.formula_columns == ()
    assert spec.formula_mask == (), "无公式列 ⇒ mask 为空"


def test_e111_rejects_row_projection() -> None:
    """无行身份 ⇒ 走行表投影必 fail closed（不会被误当动态行接入）。"""
    with pytest.raises(ENGINE.RowTableStorePayloadError, match="static_region"):
        list(ENGINE.iter_store_rows(E111.SPEC_E111, [{"id": "x"}]))


def test_e111_store_kind_is_fixed_text_with_empty_default() -> None:
    """fixed_text 形态的 per-item 缺省是空串，**不是** "[]"（blanket 默认曾致 opaque 500）。"""
    from app.services.workpaper_sync.store_item_registry import StoreKind, default_payload_for

    assert E111.SPEC_E111.store_kind is StoreKind.fixed_text
    assert E111.SPEC_E111.empty_payload == ""
    assert default_payload_for(StoreKind.fixed_text) == ""


def test_e111_has_no_rows_store_key() -> None:
    """零 `-rows` 键是它判 static_region 的第一条三元组依据。"""
    for item in E111.STORE_ITEM_IDS_E111:
        assert not item.endswith("-rows"), f"{item} 不应是 -rows 形态"
    assert not E111.STORE_ITEM_ID_PREFIX_E111.endswith("-rows")


def test_e111_keys_exist_in_host_vue_not_composable() -> None:
    """🔴 E1-11 的键在 **.vue 宿主**里而非 composable（组件明写 No composable）。

    本判据打红过一次真实的搜索范围错误：首版只在 `composables/useE1*.ts` 里 grep，
    找不到 `E1-account-commit-check-summary` 就以为键名错了 —— 实际是搜索范围漏了 `.vue`。
    """
    host = _BACKEND.parent / E111.HOST_COMPONENT_E111
    assert host.exists(), f"宿主组件缺失：{host}"
    src = host.read_text(encoding="utf-8")
    assert "No composable" in src, "本张应是无 composable 的段落表单形态"
    assert f"const ITEM_PREFIX = '{E111.STORE_ITEM_ID_PREFIX_E111}'" in src
    for item in E111.STORE_ITEM_IDS_E111:
        assert f"'{item}'" in src, f"键 {item} 不在宿主组件里"


def test_e111_naming_inconsistency_trap_is_pinned() -> None:
    """🔴 命名不一致陷阱：note/conclusion 用 `E1-commit-` 前缀，主键用 `E1-account-commit`。

    照「统一前缀派生」推演会造出 `E1-account-commit-audit-note` 这种**不存在**的键。
    """
    assert "E1-commit-audit-note" in E111.STORE_ITEM_IDS_E111
    assert "E1-commit-audit-conclusion" in E111.STORE_ITEM_IDS_E111
    # 反证：拼出来的「统一前缀」形态不在真实键集里
    for fabricated in (
        f"{E111.STORE_ITEM_ID_PREFIX_E111}-audit-note",
        f"{E111.STORE_ITEM_ID_PREFIX_E111}-audit-conclusion",
    ):
        assert fabricated not in E111.STORE_ITEM_IDS_E111, (
            f"{fabricated} 是推演造出的键，实际不存在"
        )
    host = _BACKEND.parent / E111.HOST_COMPONENT_E111
    src = host.read_text(encoding="utf-8")
    assert f"'{E111.STORE_ITEM_ID_PREFIX_E111}-audit-note'" not in src


def test_e111_cross_sheet_snapshot_is_written_by_e110() -> None:
    """E1-10 ↔ E1-11 联动键由 **E1-10 侧**写入 ⇒ 受管本张时不得把它当自身 store。"""
    comp = (
        _BACKEND.parent / "audit-platform" / "frontend" / "src" / "components"
        / "workpaper" / "composables" / "useE1AccountList.ts"
    )
    src = comp.read_text(encoding="utf-8")
    assert f"= '{E111.CROSS_SHEET_SNAPSHOT_KEY_E111}'" in src, (
        "联动键应由 useE1AccountList（E1-10 侧）声明"
    )
    assert E111.CROSS_SHEET_SNAPSHOT_KEY_E111 not in E111.STORE_ITEM_IDS_E111, (
        "联动键不得进本张的受管 store 清单（否则两方向写同一键）"
    )


def test_e111_template_table_area_is_blank(e1_first_volume) -> None:
    """模板实测：R10 有表头但 R11-19 完全空白（承诺函里的空白填写区）。"""
    ws = e1_first_volume[E111.MANAGED_SHEET_E111]
    assert ws.cell(row=E111.HEADER_ROW_E111, column=1).value == "开户银行名称"
    blank = all(
        ws.cell(row=r, column=c).value in (None, "", " ")
        for r in range(E111.FIRST_STATIC_ROW_E111, E111.LAST_STATIC_ROW_E111 + 1)
        for c in range(1, 7)
    )
    assert blank, "R11-19 应为空白 —— 若有数据则本张可能是动态行表，须重判形态"


def test_e111_static_anchors_match_template(e1_first_volume) -> None:
    """静态受管格锚点必须与模板实测内容吻合（绝对坐标，禁推演）。"""
    from openpyxl.utils import column_index_from_string

    ws = e1_first_volume[E111.MANAGED_SHEET_E111]
    for field_key, cell_ref, desc in E111.STATIC_CELL_ANCHORS_E111:
        col = "".join(ch for ch in cell_ref if ch.isalpha())
        row = int("".join(ch for ch in cell_ref if ch.isdigit()))
        val = ws.cell(row=row, column=column_index_from_string(col)).value
        assert val is not None and str(val).strip() != "", (
            f"锚点 {cell_ref}（{field_key} / {desc}）在模板里为空 —— 锚点须落在有内容的格"
        )


def test_e111_ocr_second_writer_is_declared() -> None:
    """E1-P16 的可执行落点：OCR 确认弹窗必须显式登记（受管后须在 OO 编辑态 disabled）。"""
    assert E111.OCR_DIALOG_COMPONENTS_E111, "OCR 第二写入方必须登记"


# ═══════════════════════════════════════════════════════════════════════════
# 反证：E1-9 / E1-10 不是 static_region（自省变异的另一半）
# ═══════════════════════════════════════════════════════════════════════════


def test_e109_and_e110_are_not_static_region() -> None:
    """🔴 反证：E1-9/E1-10 各只 7 公式但**有 -rows 键** ⇒ 动态行表，不得判 static_region。

    若把它们声明成 static_region，`_plan_static_writes` 会按绝对坐标直写，
    插行后坐标全错（spec Task 4 变异④）。
    """
    comp_dir = (
        _BACKEND.parent / "audit-platform" / "frontend" / "src"
        / "components" / "workpaper" / "composables"
    )
    all_src = "\n".join(
        p.read_text(encoding="utf-8", errors="ignore") for p in comp_dir.glob("useE1*.ts")
    )
    # E1-9（cert variant）走模板化键；E1-10 走字面量键
    assert "`E1-cash-count-${variant}-rows`" in all_src
    assert "E1-account-list-rows" in all_src
    # 而 E1-11 的键都不是 -rows（它们在 .vue 宿主里，由
    # test_e111_keys_exist_in_host_vue_not_composable 单独钉住）
    for item in E111.STORE_ITEM_IDS_E111:
        assert not item.endswith("-rows")

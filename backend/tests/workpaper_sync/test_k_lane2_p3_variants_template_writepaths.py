# -*- coding: utf-8 -*-
"""K 循环 lane 2 — Task 14~16：变体轴口径 + 模板层基线 + 写路径收口。

spec: k8-k9-k11-k12-k13-dedicated-composable-and-cross-cycle-hub
Task 14: K8/K9 截止性测试双 sheet 的变体轴口径裁定
Task 15: 本 lane 模板层基线（含括号缺陷分摊对账）
Task 16: 写路径与其余收口
Property: KB-P29, KB-P30, KB-P31, KB-P32, KB-P37, KB-P38, KB-P39, KB-P40

═══ 🔴 本组最贵的两个结论 ═══

**① 「同尾码双 sheet」的口径必须是「同册内」。**
跨册分组必然撞（每册都有 `审定表K{n}-1`）—— 现算跨册有 10 组同（前缀,尾码），
同册内是 **0**。HC-5 / JC-11 说的是同册内两张 sheet 共用一个尾码，K 全域为 0。

**② OCR 的两个契约面里，无 `wpId` 段的那个后端根本没注册 ⇒ 5 处是 404。**
后端只有 `/api/workpapers/{wp_id}/d4/contract-ocr`
（`backend/app/routers/wp_render_strategies/_d4_contract_ocr.py`），
没有 `/api/d4/contract-ocr`。所以 KC-19 的「统一」不是风格问题而是**修 404**。

═══ 🔴 口径差异（如实登记，不改判据方向）═══

**K6-5/K6-6 的前缀其实不同。** design 把它登记为「同型的对象变体」，但现算
`减值准备测试表（后续计量） K6-5` 与 `处置组减值测试表（后续计量） K6-6` 去括号后
前缀是「减值准备测试表」vs「处置组减值测试表」—— **不相同**。
⇒ K 的变体轴有两个子形态，判据分开写：
  - **方向变体**（K8-6/K8-7 · K9-6/K9-7）：前缀相同，差异在括号内的方向描述
  - **对象变体**（K6-5/K6-6）：前缀不同（对象名在前缀里），括号内容相同
「前缀相同」只对方向变体成立，写成通用判据会把 K6 判成不同型。
"""
from __future__ import annotations

import json
import pathlib
import re

import pytest
from openpyxl import load_workbook

from tests.workpaper_sync.k_foundation_facts import (  # noqa: E402
    BP5_INDEXES,
    DATA,
    ROOT,
    cached_text,
    derived_total_keys,
    entry_of_path,
    k_domain_files,
    strip_comments,
)

TPL_DIR = ROOT / "backend" / "wp_templates" / "K"
SLICE_PATH = DATA / "workpaper_sync_k_cycle_manifest_slice.json"

LANE2 = (8, 9, 11, 12, 13)
LANE1 = BP5_INDEXES
FOUNDATION = (10,)

#: 本 lane 的册（按文件名前缀锚定，不写死全名）
LANE2_BOOK_PREFIXES = ("K8 ", "K9 ", "K11 ", "K12 ", "K13 ")

#: 四张截止性测试 sheet 的几何（r, c, f, merged）
CUTOFF_GEOMETRY = {
    "K8-6": (44, 11, 7, 8),
    "K9-6": (44, 11, 7, 8),
    "K8-7": (44, 12, 7, 9),
    "K9-7": (44, 12, 7, 9),
}

#: 「调整分录汇总」同构指纹（与 canary 所在的 `调整分录汇总K10-3` 逐项相同）
ADJ_SUMMARY_FINGERPRINT = {"c": 10, "f": 7, "bare_if": 0, "merged": 3}


# ════════════════════════════════════════════════════════════════════════════
# 模板层口径（本文件自包含，与 foundation p2 同实现）
# ════════════════════════════════════════════════════════════════════════════
def _sheets(p: pathlib.Path) -> list[str]:
    wb = load_workbook(p, read_only=True, data_only=True)
    names = list(wb.sheetnames)
    wb.close()
    return names


def _is_bare_if(v: object) -> bool:
    return (
        isinstance(v, str)
        and re.search(r"\bIF\s*\(", v) is not None
        and re.search(r"IFERROR|IFNA|IFS\s*\(", v) is None
    )


def _bare_if_cells(p: pathlib.Path, sheet: str | None = None) -> int:
    wb = load_workbook(p, read_only=True, data_only=False)
    names = [sheet] if sheet else wb.sheetnames
    n = sum(
        1
        for sn in names
        for row in wb[sn].iter_rows()
        for c in row
        if _is_bare_if(c.value)
    )
    wb.close()
    return n


def _geometry(p: pathlib.Path, sheet: str) -> tuple[int, int, int, int]:
    """(max_row, max_col, formula_cells, merged_ranges)。"""
    wb = load_workbook(p, read_only=False, data_only=False)
    ws = wb[sheet]
    f = sum(
        1
        for row in ws.iter_rows()
        for c in row
        if isinstance(c.value, str) and c.value.startswith("=")
    )
    out = (ws.max_row, ws.max_column, f, len(ws.merged_cells.ranges))
    wb.close()
    return out


def _defined_names(p: pathlib.Path) -> tuple[int, int]:
    """(总数, 含 #REF! 数)。"""
    wb = load_workbook(p, read_only=False, data_only=False)
    total = ref_err = 0
    for dn in wb.defined_names.values():
        total += 1
        if "#REF!" in str(dn.value or ""):
            ref_err += 1
    for ws in wb.worksheets:
        local = getattr(ws, "defined_names", None)
        if not local:
            continue
        for dn in local.values():
            total += 1
            if "#REF!" in str(dn.value or ""):
                ref_err += 1
    wb.close()
    return total, ref_err


#: 🔴 KC-12 第三形态「跳跃式 SUM 多段」：SUM 参数含逗号分段
JUMP_SUM_RX = re.compile(r"SUM\(\s*[^()]*,[^()]*\)", re.I)


def _jump_sum_stats(p: pathlib.Path) -> tuple[int, set[int], set[str]]:
    """返回 (格数, 行号集合, sheet 名集合)。三个口径一次算完。"""
    wb = load_workbook(p, read_only=True, data_only=False)
    cells = 0
    rows: set[int] = set()
    sheets_: set[str] = set()
    for sn in wb.sheetnames:
        for row in wb[sn].iter_rows():
            for c in row:
                v = c.value
                if isinstance(v, str) and v.startswith("=") and JUMP_SUM_RX.search(v):
                    cells += 1
                    rows.add(c.row)
                    sheets_.add(sn)
    wb.close()
    return cells, rows, sheets_


def _ref_error_cells(p: pathlib.Path) -> int:
    wb = load_workbook(p, read_only=True, data_only=False)
    n = sum(
        1
        for sn in wb.sheetnames
        for row in wb[sn].iter_rows()
        for c in row
        if isinstance(c.value, str) and "#REF!" in c.value
    )
    wb.close()
    return n


@pytest.fixture(scope="module")
def books() -> list[pathlib.Path]:
    return sorted(p for p in TPL_DIR.rglob("*.xlsx") if not p.name.startswith("~$"))


@pytest.fixture(scope="module")
def lane2_books(books: list[pathlib.Path]) -> list[pathlib.Path]:
    out = [p for p in books if p.name.startswith(LANE2_BOOK_PREFIXES)]
    assert len(out) == 5, f"本 lane 册期望 5 本，实得 {[p.name for p in out]}"
    return out


@pytest.fixture(scope="module")
def k_files() -> list[pathlib.Path]:
    return k_domain_files()


@pytest.fixture(scope="module")
def slice_doc() -> dict:
    return json.loads(SLICE_PATH.read_text(encoding="utf-8"))


def _entry_of_book(p: pathlib.Path) -> int | None:
    m = re.match(r"^K(\d+)\s", p.name)
    return int(m.group(1)) if m else None


def _per_entry_hits(
    files: list[pathlib.Path], rx: re.Pattern[str]
) -> dict[int | None, int]:
    out: dict[int | None, int] = {}
    for p in files:
        n = len(rx.findall(strip_comments(cached_text(p))))
        if n:
            k = entry_of_path(p)
            out[k] = out.get(k, 0) + n
    return out


# ════════════════════════════════════════════════════════════════════════════
# Task 14 / KB-P29~P30, P32：变体轴口径裁定
# ════════════════════════════════════════════════════════════════════════════
class TestKBP29CutoffSheetGeometry:
    """四张截止性测试 sheet 的尾码与几何。"""

    def test_four_cutoff_sheets_exist(self, books: list[pathlib.Path]) -> None:
        found = {
            sn
            for p in books
            for sn in _sheets(p)
            if re.search(r"K[89]-[67]$", sn)
        }
        assert len(found) == 4, f"截止性 sheet 期望 4 张，实得 {sorted(found)}"

    @pytest.mark.parametrize("code", sorted(CUTOFF_GEOMETRY))
    def test_geometry_matches_design(
        self, code: str, books: list[pathlib.Path]
    ) -> None:
        """r=44 恒定；c 与 merged 按方向分 11/8 与 12/9。"""
        hit = [
            (p, sn) for p in books for sn in _sheets(p) if sn.endswith(code)
        ]
        assert len(hit) == 1, f"{code} 应唯一命中，实得 {len(hit)}"
        p, sn = hit[0]
        assert _geometry(p, sn) == CUTOFF_GEOMETRY[code], (
            f"{sn} 几何 {_geometry(p, sn)} != 期望 {CUTOFF_GEOMETRY[code]}"
        )

    def test_row_count_is_identical_across_all_four(self) -> None:
        """四张 r 全 44 ⇒ 差异只在列数（方向多一列）。"""
        assert {g[0] for g in CUTOFF_GEOMETRY.values()} == {44}

    def test_direction_adds_exactly_one_column_and_one_merge(self) -> None:
        """🔴 `-7`（原始→记账）比 `-6`（记账→原始）多 1 列 + 1 合并区。"""
        for n in (8, 9):
            c6 = CUTOFF_GEOMETRY[f"K{n}-6"]
            c7 = CUTOFF_GEOMETRY[f"K{n}-7"]
            assert c7[1] - c6[1] == 1, f"K{n}: 列差 {c7[1] - c6[1]}"
            assert c7[3] - c6[3] == 1, f"K{n}: 合并区差 {c7[3] - c6[3]}"
            assert c6[2] == c7[2] == 7, "公式格数应两侧同为 7"

    def test_k8_and_k9_are_geometrically_identical(self) -> None:
        """K8 与 K9 的同尾码 sheet 几何完全一致（同构双册）。"""
        assert CUTOFF_GEOMETRY["K8-6"] == CUTOFF_GEOMETRY["K9-6"]
        assert CUTOFF_GEOMETRY["K8-7"] == CUTOFF_GEOMETRY["K9-7"]


class TestKBP29SameSuffixIsZeroInKDomain:
    """🔴 HC-5 / JC-11 的「同尾码双 sheet」在 K 全域为 0（照抄会误判）。"""

    def test_same_suffix_within_a_book_is_zero(
        self, books: list[pathlib.Path]
    ) -> None:
        """🔴 口径必须是**同册内** —— 跨册分组必然撞。"""
        dupes: dict[str, dict] = {}
        for p in books:
            groups: dict[tuple[str, str], list[str]] = {}
            for sn in _sheets(p):
                m = re.match(r"^(.*?)\s*(K\d+)-(\d+)$", sn)
                if m:
                    key = (m.group(1).strip(), m.group(3))
                    groups.setdefault(key, []).append(sn)
            d = {k: v for k, v in groups.items() if len(v) > 1}
            if d:
                dupes[p.name] = d
        assert dupes == {}, f"同册内出现同（前缀,尾码）：{dupes}"

    def test_cross_book_grouping_would_falsely_report_many(
        self, books: list[pathlib.Path]
    ) -> None:
        """🔴 变异反证：按**跨册**分组会得到一堆假命中 ⇒ 证明口径选择要紧。"""
        groups: dict[tuple[str, str], list[str]] = {}
        for p in books:
            for sn in _sheets(p):
                m = re.match(r"^(.*?)\s*(K\d+)-(\d+)$", sn)
                if m:
                    groups.setdefault(
                        (m.group(1).strip(), m.group(3)), []
                    ).append(sn)
        cross = {k: v for k, v in groups.items() if len(v) > 1}
        assert len(cross) >= 5, (
            f"跨册分组只得 {len(cross)} 组 ⇒ 「口径要紧」的反证不成立"
        )
        # 典型假命中：每册都有的「审定表 -1」
        assert any(k[1] == "1" and len(v) >= 8 for k, v in cross.items()), (
            "找不到「审定表 -1」这类跨册重名 ⇒ 反证样本变了"
        )

    def test_the_four_cutoff_sheets_have_distinct_suffixes(self) -> None:
        """四张 sheet 的尾码是 6/7 两两不同（不是同尾码）。"""
        for n in (8, 9):
            assert {f"K{n}-6", f"K{n}-7"} <= set(CUTOFF_GEOMETRY)
        suffixes = {c.rsplit("-", 1)[1] for c in CUTOFF_GEOMETRY}
        assert suffixes == {"6", "7"}


class TestKBP30VariantAxisIsAdjacencyNotSameSuffix:
    """🔴 K 的变体轴判据 = 「尾码相邻 + 语义成对」，**不是「尾码相同」**。"""

    def test_direction_variants_share_the_prefix(
        self, books: list[pathlib.Path]
    ) -> None:
        """方向变体：前缀相同（差异在括号内），尾码相邻。"""
        found: list[tuple[str, str]] = []
        for p in books:
            pref: dict[str, list[tuple[int, str]]] = {}
            for sn in _sheets(p):
                m = re.match(r"^(.*?)\s*(K\d+)-(\d+)$", sn)
                if not m:
                    continue
                base = re.sub(r"[（(].*?[）)]", "", m.group(1)).strip()
                pref.setdefault(base, []).append((int(m.group(3)), sn))
            for base, items in pref.items():
                items.sort()
                for a, b in zip(items, items[1:]):
                    if b[0] - a[0] == 1:
                        found.append((a[1], b[1]))
        assert len(found) == 2, (
            f"方向变体对数期望 2（K8/K9 各 1），实得 {found}"
        )
        for a, b in found:
            assert "截止性测试" in a and "截止性测试" in b

    def test_object_variant_k6_has_a_different_prefix(
        self, books: list[pathlib.Path]
    ) -> None:
        """🔴 口径差异：K6-5/K6-6 的前缀**不相同**（对象名在前缀里）。

        design 把它登记为「同型的对象变体」—— 同型指的是「尾码相邻 + 语义成对」
        这个轴，**不是**「前缀相同」。把「前缀相同」写成通用判据会把 K6 判成不同型。
        """
        k6 = next(p for p in books if p.name.startswith("K6 "))
        names = {
            sn: re.sub(r"[（(].*?[）)]", "", sn.rsplit("K6-", 1)[0]).strip()
            for sn in _sheets(k6)
            if re.search(r"K6-[56]$", sn)
        }
        assert len(names) == 2, f"K6-5/K6-6 应各 1 张，实得 {names}"
        prefixes = set(names.values())
        assert len(prefixes) == 2, (
            f"K6-5/K6-6 前缀相同（{prefixes}）⇒ 本口径差异登记须撤"
        )

    def test_object_variant_still_satisfies_the_k_axis(
        self, books: list[pathlib.Path]
    ) -> None:
        """两侧都验：K6-5/K6-6 满足「尾码相邻 + 括号内容相同」。"""
        k6 = next(p for p in books if p.name.startswith("K6 "))
        pairs = sorted(
            sn for sn in _sheets(k6) if re.search(r"K6-[56]$", sn)
        )
        assert len(pairs) == 2
        sufs = sorted(int(sn.rsplit("-", 1)[1]) for sn in pairs)
        assert sufs[1] - sufs[0] == 1, "尾码不相邻 ⇒ 不在同一变体轴"
        parens = {
            tuple(re.findall(r"[（(](.*?)[）)]", sn)) for sn in pairs
        }
        assert len(parens) == 1, f"括号内容不同：{parens} ⇒ 不是对象变体"

    def test_k6_pair_is_in_lane1_not_this_lane(
        self, books: list[pathlib.Path]
    ) -> None:
        """K6 归 lane 1 —— 本 lane 只交付方向变体那两对。"""
        k6 = next(p for p in books if p.name.startswith("K6 "))
        assert _entry_of_book(k6) == 6
        assert 6 in LANE1 and 6 not in LANE2


class TestKBP32CutoffSamplingIsNotAChecklistWritePath:
    """`sampling/cutoff-test` 是取数端点，不是 checklist 写路径。"""

    def test_lane2_has_exactly_two_call_sites(
        self, k_files: list[pathlib.Path]
    ) -> None:
        """🔴 本 lane 2 个文件（`useK8Cutoff.ts` / `useK9Cutoff.ts`）。"""
        rx = re.compile(r"sampling/cutoff-test")
        hits = {
            p.name
            for p in k_files
            if entry_of_path(p) in LANE2 and rx.search(strip_comments(cached_text(p)))
        }
        assert hits == {"useK8Cutoff.ts", "useK9Cutoff.ts"}, hits

    def test_it_is_a_read_endpoint_not_a_write(
        self, k_files: list[pathlib.Path]
    ) -> None:
        """🔴 判据：调用点不带 `checklist-responses`，且落点是本地 ref 不是 onSave。"""
        for fname in ("useK8Cutoff.ts", "useK9Cutoff.ts"):
            p = next(x for x in k_files if x.name == fname)
            src = strip_comments(cached_text(p))
            m = re.search(
                r"sampling/cutoff-test[\s\S]{0,600}", src
            )
            assert m is not None
            seg = m.group(0)
            assert "checklist-responses" not in seg, (
                f"{fname} 的 cutoff-test 段内出现 checklist-responses ⇒ 它是写路径"
            )
            assert re.search(r"samples\.value\s*=", seg), (
                f"{fname} 的 cutoff-test 结果未落到本地 samples ref"
            )

    def test_backend_route_exists(self) -> None:
        """两侧都验：后端真有这个取数路由（不是猜的端点）。"""
        hits = [
            p.relative_to(ROOT).as_posix()
            for p in (ROOT / "backend" / "app" / "routers").rglob("*.py")
            if "sampling/cutoff-test" in p.read_text(encoding="utf-8", errors="ignore")
        ]
        assert hits, "后端找不到 sampling/cutoff-test ⇒ 端点是猜的"


# ════════════════════════════════════════════════════════════════════════════
# Task 15 / KB-P37, KB-P31：模板层基线
# ════════════════════════════════════════════════════════════════════════════
#: 本 lane 逐册规模（sheets, bare_if）
LANE2_BOOK_SCALE = {
    8: (12, 136),
    9: (12, 175),
    11: (7, 39),
    12: (9, 37),
    13: (9, 37),
}
#: 13 entry 册裸 IF 合计（排除 K0 函证册）
ENTRY_BOOKS_BARE_IF_TOTAL = 715


class TestKBP37Lane2TemplateScale:
    """sheets 49 · 裸 IF 424（占 13 entry 册 715 的 59%）。"""

    @pytest.mark.parametrize("n", sorted(LANE2_BOOK_SCALE))
    def test_per_book_scale(self, n: int, books: list[pathlib.Path]) -> None:
        p = next(x for x in books if _entry_of_book(x) == n)
        exp_sheets, exp_if = LANE2_BOOK_SCALE[n]
        assert len(_sheets(p)) == exp_sheets, (
            f"K{n} sheets 期望 {exp_sheets}，实得 {len(_sheets(p))}"
        )
        assert _bare_if_cells(p) == exp_if, (
            f"K{n} 裸 IF 期望 {exp_if}，实得 {_bare_if_cells(p)}"
        )

    def test_sheets_total_is_49(self, lane2_books: list[pathlib.Path]) -> None:
        total = sum(len(_sheets(p)) for p in lane2_books)
        assert total == 49, f"本 lane sheets 期望 49，实得 {total}"
        assert total == sum(s for s, _ in LANE2_BOOK_SCALE.values())

    def test_k11_has_the_fewest_sheets_in_whole_k(
        self, books: list[pathlib.Path]
    ) -> None:
        """🔴 K11 只 7 张是全 K 最少（两侧都验：含非本 lane 的册）。"""
        entry_books = [p for p in books if _entry_of_book(p) != 0]
        counts = {_entry_of_book(p): len(_sheets(p)) for p in entry_books}
        assert counts[11] == 7
        assert counts[11] == min(counts.values()), (
            f"K11 不是最少：{dict(sorted(counts.items()))}"
        )
        assert sum(1 for v in counts.values() if v == 7) == 1, "最少值不唯一"

    def test_bare_if_total_is_424(self, lane2_books: list[pathlib.Path]) -> None:
        total = sum(_bare_if_cells(p) for p in lane2_books)
        assert total == 424, f"本 lane 裸 IF 期望 424，实得 {total}"

    def test_bare_if_share_is_59_percent(
        self, books: list[pathlib.Path]
    ) -> None:
        """🔴 占 13 entry 册 715 的 59%（三份 spec 最高）。"""
        entry_books = [p for p in books if _entry_of_book(p) != 0]
        assert len(entry_books) == 13, f"entry 册期望 13 本，实得 {len(entry_books)}"
        grand = sum(_bare_if_cells(p) for p in entry_books)
        assert grand == ENTRY_BOOKS_BARE_IF_TOTAL, (
            f"13 entry 册裸 IF 期望 {ENTRY_BOOKS_BARE_IF_TOTAL}，实得 {grand}"
        )
        lane2 = sum(
            _bare_if_cells(p) for p in entry_books if _entry_of_book(p) in LANE2
        )
        share = lane2 / grand
        assert 0.585 <= share <= 0.60, f"占比 {share:.3%} 不在 59% 附近"
        lane1 = sum(
            _bare_if_cells(p) for p in entry_books if _entry_of_book(p) in LANE1
        )
        assert lane2 > lane1, "本 lane 未占最高 ⇒ design 表述须改"

    def test_k9_then_k8_are_the_top_two(
        self, books: list[pathlib.Path]
    ) -> None:
        """K9 175 与 K8 136 是全 K 前二。"""
        entry_books = [p for p in books if _entry_of_book(p) != 0]
        ranked = sorted(
            ((_bare_if_cells(p), _entry_of_book(p)) for p in entry_books),
            reverse=True,
        )
        assert [e for _c, e in ranked[:2]] == [9, 8], f"前二 {ranked[:3]}"


class TestKBP37StructuralZerosInLane2:
    """definedName 5 册全 0 · `#REF!` 0 格 · 跳跃式 footer 0 处。"""

    def test_defined_names_are_zero(self, lane2_books: list[pathlib.Path]) -> None:
        for p in lane2_books:
            total, ref_err = _defined_names(p)
            assert total == 0, f"{p.name}: definedName {total} != 0"
            assert ref_err == 0

    def test_ref_error_cells_are_zero(
        self, lane2_books: list[pathlib.Path]
    ) -> None:
        for p in lane2_books:
            assert _ref_error_cells(p) == 0, f"{p.name} 有 #REF! 格"

    def test_lane1_has_the_ref_errors_for_contrast(
        self, books: list[pathlib.Path]
    ) -> None:
        """🔴 反向对照：#REF! 在 lane 1 的 K1/K2 册非 0（本 lane 的 0 有意义）。"""
        lane1_total = sum(
            _defined_names(p)[1]
            for p in books
            if _entry_of_book(p) in LANE1
        )
        assert lane1_total > 0, (
            "lane 1 也没有 #REF! ⇒ 本 lane「0」不构成对照"
        )

    def test_no_jump_style_footer(self, lane2_books: list[pathlib.Path]) -> None:
        """跳跃式 footer 0 处 —— 三个口径同时为 0。

        🔴 形态是 **KC-12 的第三个 footer 形态「跳跃式 SUM 多段」**：
        `=SUM(B9,B10)` / `=SUM(B16,B26,B37,B45)` / `=SUM(B25:B26,B29:B30)`
        这种 SUM 参数含逗号分段的公式 —— **不是**「合计行上方有空行」那种版式特征
        （按后者扫本 lane 会得 9 处假命中）。
        """
        for p in lane2_books:
            cells, rows, sheets_ = _jump_sum_stats(p)
            assert (cells, len(rows), len(sheets_)) == (0, 0, 0), (
                f"{p.name} 有跳跃式 SUM：格 {cells} 行 {sorted(rows)} sheet {sheets_}"
            )

    def test_jump_style_is_concentrated_in_lane1(
        self, books: list[pathlib.Path]
    ) -> None:
        """🔴 反向对照：全 K 的跳跃式 SUM 全在 lane 1 的 K1 与 K6（本 lane 0 有意义）。"""
        per: dict[int | None, int] = {}
        for p in books:
            cells, _rows, _sheets = _jump_sum_stats(p)
            if cells:
                per[_entry_of_book(p)] = cells
        assert set(per) == {1, 6}, f"跳跃式分布 {per} ⇒ 对照对象变了"
        assert sum(v for k, v in per.items() if k in LANE2) == 0

    def test_design_five_is_a_mixed_caliber(
        self, books: list[pathlib.Path]
    ) -> None:
        """🔴 口径差异登记：design 的「跳跃式 5」混用了行口径与 sheet 口径。

        现算全 K：**35 格 / 6 行组 / 3 sheet**。
        design 的 5 = `审定表K1-1` 按**行**算 3 + `坏账准备测算K1-8` 按 **sheet** 算 1
        + `初始确认检查表K6-4` 按 **sheet** 算 1（该 sheet 实为 2 行）。
        ⇒ 三个口径任选其一都自洽，但 5 不属于任何单一口径。
        本 lane 的「0」在三个口径下都成立，所以本判据不受该差异影响。
        """
        cells = rows = 0
        sheets_: set[str] = set()
        for p in books:
            c, r, s = _jump_sum_stats(p)
            cells += c
            rows += len(r)
            sheets_ |= {f"{p.name}/{x}" for x in s}
        assert cells == 35, f"格口径期望 35，实得 {cells}"
        assert rows == 6, f"行组口径期望 6，实得 {rows}"
        assert len(sheets_) == 3, f"sheet 口径期望 3，实得 {sorted(sheets_)}"
        assert 5 not in (cells, rows, len(sheets_)), (
            "design 的 5 已能对应某个单一口径 ⇒ 本差异登记可撤"
        )

    def test_gt_custom_split_is_k12_k13_only(
        self, lane2_books: list[pathlib.Path]
    ) -> None:
        """🔴 hidden `GT_Custom`：K12/K13 **有**、K8/K9/K11 **无**。"""
        has = {
            _entry_of_book(p): ("GT_Custom" in _sheets(p)) for p in lane2_books
        }
        assert has == {8: False, 9: False, 11: False, 12: True, 13: True}, has


class TestKBP31SheetNameDefectsAreApportioned:
    """🔴 括号缺陷分摊对账：本 lane 2 + lane 1 的 3 == 5（KC-10）。"""

    @staticmethod
    def _bracket_defects(p: pathlib.Path) -> list[str]:
        out: list[str] = []
        for sn in _sheets(p):
            ho, hc = sn.count("("), sn.count(")")
            fo, fc = sn.count("（"), sn.count("）")
            mixed = (ho + hc > 0) and (fo + fc > 0)
            unpaired = (ho != hc) or (fo != fc)
            if mixed or unpaired:
                out.append(sn)
        return out

    def test_lane2_has_two_bracket_defects(
        self, lane2_books: list[pathlib.Path]
    ) -> None:
        hits = [
            (p.name, sn) for p in lane2_books for sn in self._bracket_defects(p)
        ]
        assert len(hits) == 2, f"本 lane 括号缺陷期望 2，实得 {hits}"
        codes = {sn[-4:] for _n, sn in hits}
        assert codes == {"K8-6", "K9-6"}, f"命中的 sheet 码 {codes}"

    def test_two_plus_three_equals_five(self, books: list[pathlib.Path]) -> None:
        """🔴 全 K 5 处 = 本 lane 2 + lane 1 的 3（逐侧现算）。"""
        per_lane = {"lane1": 0, "lane2": 0, "other": 0}
        for p in books:
            n = len(self._bracket_defects(p))
            if not n:
                continue
            e = _entry_of_book(p)
            key = "lane2" if e in LANE2 else "lane1" if e in LANE1 else "other"
            per_lane[key] += n
        assert per_lane["lane2"] == 2, per_lane
        assert per_lane["lane1"] == 3, per_lane
        assert per_lane["other"] == 0, f"K0/K10 也有括号缺陷：{per_lane}"
        assert sum(per_lane.values()) == 5

    def test_half_width_space_in_name_is_one_in_this_lane(
        self, lane2_books: list[pathlib.Path]
    ) -> None:
        """🔴 名中半角空格本 lane 只 1 处（`实质性程序表 K11A`）。"""
        hits = [
            (p.name, sn) for p in lane2_books for sn in _sheets(p) if " " in sn
        ]
        assert len(hits) == 1, f"本 lane 半角空格期望 1，实得 {hits}"
        assert hits[0][1] == "实质性程序表 K11A", hits

    def test_half_width_space_is_concentrated_in_lane1(
        self, books: list[pathlib.Path]
    ) -> None:
        """🔴 与 lane 1 对账：绝大多数在 K4/K5/K6（那三册整体带空格）。"""
        per_entry: dict[int | None, int] = {}
        for p in books:
            n = sum(1 for sn in _sheets(p) if " " in sn)
            if n:
                per_entry[_entry_of_book(p)] = n
        lane2 = sum(v for k, v in per_entry.items() if k in LANE2)
        lane1 = sum(v for k, v in per_entry.items() if k in LANE1)
        assert lane2 == 1, f"本 lane {lane2}（分布 {per_entry}）"
        assert lane1 >= 18, f"lane 1 期望 ≥18，实得 {lane1}"
        assert {4, 5, 6} <= set(per_entry), "K4/K5/K6 应都命中"


class TestKBP37AdjustmentSummaryIsIsomorphicToCanary:
    """🔴 5 张「调整分录汇总」与 canary 的 `调整分录汇总K10-3` 完全同构。

    这是 canary 判据能外推到本 lane 的**直接依据** —— 否则拿 K10 验出来的
    结论对 K8/K9/K11/K12/K13 一句话都说不上。
    """

    @staticmethod
    def _fingerprint(p: pathlib.Path, sn: str) -> dict:
        _r, c, f, merged = _geometry(p, sn)
        return {"c": c, "f": f, "bare_if": _bare_if_cells(p, sn), "merged": merged}

    def test_canary_sheet_fingerprint(self, books: list[pathlib.Path]) -> None:
        """先立 canary 侧指纹（c=10 f=7 bareIF=0 merged=3）。"""
        k10 = next(p for p in books if _entry_of_book(p) == 10)
        sn = next(s for s in _sheets(k10) if s == "调整分录汇总K10-3")
        assert self._fingerprint(k10, sn) == ADJ_SUMMARY_FINGERPRINT

    @pytest.mark.parametrize("n", sorted(LANE2_BOOK_SCALE))
    def test_each_lane2_sheet_matches_the_canary(
        self, n: int, books: list[pathlib.Path]
    ) -> None:
        p = next(x for x in books if _entry_of_book(x) == n)
        cands = [s for s in _sheets(p) if "调整分录汇总" in s]
        assert len(cands) == 1, f"K{n} 的调整分录汇总应唯一，实得 {cands}"
        assert self._fingerprint(p, cands[0]) == ADJ_SUMMARY_FINGERPRINT, (
            f"K{n} {cands[0]} 指纹 {self._fingerprint(p, cands[0])}"
            f" != canary {ADJ_SUMMARY_FINGERPRINT}"
        )

    def test_row_count_differs_but_is_not_part_of_the_fingerprint(
        self, books: list[pathlib.Path]
    ) -> None:
        """🔴 行数**不同**（21~25）且刻意不进指纹 —— 它是动态行上限不是结构。"""
        rows: dict[int | None, int] = {}
        for p in books:
            for sn in _sheets(p):
                if "调整分录汇总" in sn:
                    rows[_entry_of_book(p)] = _geometry(p, sn)[0]
        assert len(set(rows.values())) > 1, (
            f"所有册行数相同（{set(rows.values())}）⇒ 本澄清可撤"
        )
        assert "r" not in ADJ_SUMMARY_FINGERPRINT


class TestKBP37AdjudicationComplexity:
    """审定表复杂度：K9-1 f=223 / K11-1 f=190（全 K 第二、第三）。"""

    def test_k9_and_k11_formula_counts(self, books: list[pathlib.Path]) -> None:
        expect = {9: 223, 11: 190}
        for n, exp in expect.items():
            p = next(x for x in books if _entry_of_book(x) == n)
            sn = next(s for s in _sheets(p) if s.endswith(f"K{n}-1"))
            f = _geometry(p, sn)[2]
            assert f == exp, f"K{n}-1 公式格期望 {exp}，实得 {f}"

    def test_k1_is_first_and_k9_k11_follow(
        self, books: list[pathlib.Path]
    ) -> None:
        """🔴 第一是 lane 1 的 K1-1 f=547（本 lane 不是最复杂的）。"""
        ranked: list[tuple[int, int | None]] = []
        for p in books:
            e = _entry_of_book(p)
            if e == 0:
                continue
            for sn in _sheets(p):
                if re.search(rf"K{e}-1$", sn):
                    ranked.append((_geometry(p, sn)[2], e))
        ranked.sort(reverse=True)
        assert [e for _f, e in ranked[:3]] == [1, 9, 11], f"前三 {ranked[:4]}"
        assert ranked[0][0] == 547, f"K1-1 期望 547，实得 {ranked[0][0]}"
        assert 1 in LANE1, "K1 应属 lane 1"


# ════════════════════════════════════════════════════════════════════════════
# Task 16 / KB-P38~P40：写路径与其余收口
# ════════════════════════════════════════════════════════════════════════════
PUBLISH_RX = re.compile(r"audit-determination/publish-to-tb")
CHECKLIST_RX = re.compile(r"checklist-responses")
CENTRAL_SYNC_RX = re.compile(r"useAdjustmentCentralSync")
OCR_BARE_RX = re.compile(r"/api/d4/contract-ocr")
OCR_WPID_RX = re.compile(r"/api/workpapers/\$\{[^}]*\}/d4/contract-ocr")

#: TB 发布门的 5 个宿主（design 点名）
LANE2_PUBLISH_HOSTS = {
    "useK8Adjudication.ts",
    "useK9Adjudication.ts",
    "useK11Adjudication.ts",
    "K12TabAdjudication.vue",
    "K13TabAdjudication.vue",
}


class TestKBP38TbPublishGate:
    """TB 发布门 5 处 / 5 条 entry；5 + 8 + 1 == 14。"""

    def test_lane2_hosts_match_design(self, k_files: list[pathlib.Path]) -> None:
        hosts = {
            p.name
            for p in k_files
            if entry_of_path(p) in LANE2
            and PUBLISH_RX.search(strip_comments(cached_text(p)))
        }
        assert hosts == LANE2_PUBLISH_HOSTS, (
            f"多 {sorted(hosts - LANE2_PUBLISH_HOSTS)}，"
            f"缺 {sorted(LANE2_PUBLISH_HOSTS - hosts)}"
        )

    def test_one_gate_per_entry(self, k_files: list[pathlib.Path]) -> None:
        """5 条 entry 各恰 1 处（不重复发布门）。"""
        per = _per_entry_hits(
            [p for p in k_files if entry_of_path(p) in LANE2], PUBLISH_RX
        )
        assert per == {n: 1 for n in LANE2}, per

    def test_five_plus_eight_plus_one_equals_14(
        self, k_files: list[pathlib.Path]
    ) -> None:
        """🔴 三份 spec 分摊对账：本 lane 5 + lane 1 的 8 + foundation 1 == 14。"""
        per = _per_entry_hits(k_files, PUBLISH_RX)
        l2 = sum(v for k, v in per.items() if k in LANE2)
        l1 = sum(v for k, v in per.items() if k in LANE1)
        f0 = sum(v for k, v in per.items() if k in FOUNDATION)
        assert (l2, l1, f0) == (5, 8, 1), f"分摊 {(l2, l1, f0)}（分布 {per}）"
        assert l2 + l1 + f0 == 14
        assert sum(per.values()) == 14, "有发布门落在三份 spec 之外"

    def test_the_host_layer_is_mixed_composable_and_vue(self) -> None:
        """🔴 本 lane 的 5 个宿主**跨两层**（3 composable + 2 Vue）⇒ 判据不能只扫一层。"""
        ts = {h for h in LANE2_PUBLISH_HOSTS if h.endswith(".ts")}
        vue = {h for h in LANE2_PUBLISH_HOSTS if h.endswith(".vue")}
        assert len(ts) == 3 and len(vue) == 2, (ts, vue)


class TestKBP38ChecklistResponsesCarrier:
    """`checklist-responses` 载体：5 个 FormData 各 2 处 + `useK11Detail.ts` 1 处。"""

    def test_formdata_composables_have_two_each(
        self, k_files: list[pathlib.Path]
    ) -> None:
        for n in LANE2:
            p = next(
                (x for x in k_files if x.name == f"useK{n}FormData.ts"), None
            )
            assert p is not None, f"useK{n}FormData.ts 不存在"
            c = len(CHECKLIST_RX.findall(strip_comments(cached_text(p))))
            assert c == 2, f"useK{n}FormData.ts 期望 2 处，实得 {c}"

    def test_k11_detail_adds_one_more(self, k_files: list[pathlib.Path]) -> None:
        """🔴 K11 多一处（`useK11Detail.ts`）—— 本 lane 唯一的例外。"""
        p = next(x for x in k_files if x.name == "useK11Detail.ts")
        assert len(CHECKLIST_RX.findall(strip_comments(cached_text(p)))) == 1

    def test_lane2_total_is_eleven(self, k_files: list[pathlib.Path]) -> None:
        """5 × 2 + 1 == 11。"""
        per = _per_entry_hits(
            [p for p in k_files if entry_of_path(p) in LANE2], CHECKLIST_RX
        )
        assert sum(per.values()) == 11, per
        assert per[11] == 3, f"K11 期望 3（FormData 2 + Detail 1），实得 {per.get(11)}"

    def test_all_belong_to_the_bare_endpoint_family(
        self, k_files: list[pathlib.Path]
    ) -> None:
        """🔴 全属 `formdata_composable_bare_endpoint` 族（裸端点非 store 投影）。"""
        for p in k_files:
            if entry_of_path(p) not in LANE2:
                continue
            src = strip_comments(cached_text(p))
            if not CHECKLIST_RX.search(src):
                continue
            assert re.search(r"(?:http|api)\.(?:get|post|put)", src), (
                f"{p.name} 有 checklist-responses 但无裸 HTTP 调用 ⇒ 族判定变了"
            )


class TestKBP38DisclosureDualVariant:
    """披露双变体 10 处；10 + 14 + 2 == 26。"""

    def test_lane2_has_ten(self, k_files: list[pathlib.Path]) -> None:
        hits = [
            p.name
            for p in k_files
            if entry_of_path(p) in LANE2
            and re.search(r"Disclosure(?:Listed|Soe)\.vue$", p.name)
        ]
        assert len(hits) == 10, f"本 lane 披露双变体期望 10，实得 {sorted(hits)}"

    def test_every_entry_has_both_variants(
        self, k_files: list[pathlib.Path]
    ) -> None:
        """5 条 entry 各有 Listed + Soe 两个（成对，无落单）。"""
        for n in LANE2:
            got = {
                m.group(1)
                for p in k_files
                if entry_of_path(p) == n
                for m in [re.search(r"Disclosure(Listed|Soe)\.vue$", p.name)]
                if m
            }
            assert got == {"Listed", "Soe"}, f"K{n} 变体 {got}"

    def test_ten_plus_fourteen_plus_two_equals_26(
        self, k_files: list[pathlib.Path]
    ) -> None:
        """🔴 三份 spec 分摊对账。"""
        rx = re.compile(r"Disclosure(?:Listed|Soe)\.vue$")
        per: dict[str, int] = {"lane1": 0, "lane2": 0, "foundation": 0, "other": 0}
        for p in k_files:
            if not rx.search(p.name):
                continue
            e = entry_of_path(p)
            key = (
                "lane2"
                if e in LANE2
                else "lane1"
                if e in LANE1
                else "foundation"
                if e in FOUNDATION
                else "other"
            )
            per[key] += 1
        assert (per["lane2"], per["lane1"], per["foundation"]) == (10, 14, 2), per
        assert per["other"] == 0, f"有披露组件落在三份 spec 之外：{per}"
        assert sum(per.values()) == 26


class TestKBP39SecondWriterIsCentralSync:
    """`useAdjustmentCentralSync` 5 条各 3 处 ⇒ roundtrip 的第二写入方。"""

    def test_each_entry_has_three_references(
        self, k_files: list[pathlib.Path]
    ) -> None:
        per = _per_entry_hits(
            [p for p in k_files if entry_of_path(p) in LANE2], CENTRAL_SYNC_RX
        )
        assert per == {n: 3 for n in LANE2}, per

    def test_host_is_the_adjustment_tab(self, k_files: list[pathlib.Path]) -> None:
        """宿主是 `K{n}TabAdjustment.vue`（调整分录页，不是审定页）。"""
        for n in LANE2:
            hosts = {
                p.name
                for p in k_files
                if entry_of_path(p) == n
                and CENTRAL_SYNC_RX.search(strip_comments(cached_text(p)))
            }
            assert hosts == {f"K{n}TabAdjustment.vue"}, f"K{n} 宿主 {hosts}"

    def test_it_is_a_distinct_writer_from_the_tb_gate(
        self, k_files: list[pathlib.Path]
    ) -> None:
        """🔴 判据核心：第二写入方与 TB 发布门**不同宿主** ⇒ roundtrip 要算两方。"""
        for n in LANE2:
            sync_hosts = {
                p.name
                for p in k_files
                if entry_of_path(p) == n
                and CENTRAL_SYNC_RX.search(strip_comments(cached_text(p)))
            }
            gate_hosts = {
                p.name
                for p in k_files
                if entry_of_path(p) == n
                and PUBLISH_RX.search(strip_comments(cached_text(p)))
            }
            assert sync_hosts & gate_hosts == set(), (
                f"K{n} 两个写入方同宿主 {sync_hosts & gate_hosts} ⇒ 「第二写入方」表述须改"
            )

    def test_all_thirteen_entries_share_this_pattern(
        self, k_files: list[pathlib.Path]
    ) -> None:
        """两侧都验：13 条 entry 都是各 3 处（本 lane 不是特例）。"""
        per = _per_entry_hits(k_files, CENTRAL_SYNC_RX)
        assert len(per) == 13, f"命中 entry 数 {len(per)}"
        assert set(per.values()) == {3}, f"有 entry 不是 3 处：{per}"


class TestKBP40OcrTwoContractFacesUnified:
    """🔴 KC-19：OCR 的两个契约面在本 lane 统一为一种。"""

    def test_bare_form_has_no_backend_route(self) -> None:
        """🔴 修复依据：无 `wpId` 段的形态后端**根本没注册** ⇒ 那 5 处是 404。"""
        routers = ROOT / "backend" / "app" / "routers"
        declared: list[str] = []
        for p in routers.rglob("*.py"):
            src = p.read_text(encoding="utf-8", errors="ignore")
            declared += re.findall(
                r"""@router\.post\(\s*['"]([^'"]*d4/contract-ocr)['"]""", src
            )
        assert declared, "后端找不到任何 d4/contract-ocr 路由"
        assert all("{wp_id}" in d for d in declared), (
            f"存在不带 wp_id 的注册 {declared} ⇒ 404 的判定须复核"
        )
        assert not any(d == "/api/d4/contract-ocr" for d in declared)

    def test_lane2_has_zero_bare_form(self, k_files: list[pathlib.Path]) -> None:
        """本 lane 的裸形态已清零。"""
        residue = [
            p.name
            for p in k_files
            if entry_of_path(p) in LANE2
            and OCR_BARE_RX.search(strip_comments(cached_text(p)))
        ]
        assert residue == [], f"仍有 404 形态：{residue}"

    def test_lane2_all_seven_use_the_wpid_form(
        self, k_files: list[pathlib.Path]
    ) -> None:
        """本 lane 7 处全用含 `wpId` 段的形态（统一为一种）。"""
        per = _per_entry_hits(
            [p for p in k_files if entry_of_path(p) in LANE2], OCR_WPID_RX
        )
        assert sum(per.values()) == 7, f"本 lane 含 wpId 期望 7 处，实得 {per}"
        assert set(per) == {8, 9, 12, 13}, (
            f"命中 entry {sorted(per)} —— K11 无 OCR 是实况"
        )

    def test_lane1_bare_form_remains_for_its_own_spec(
        self, k_files: list[pathlib.Path]
    ) -> None:
        """🔴 反向对照：lane 1 的 2 处裸形态**还在**（归 lane 1 spec 处置）。

        这条是「本 lane 没有越界改」的正向证明。
        """
        lane1 = [
            p.name
            for p in k_files
            if entry_of_path(p) in LANE1
            and OCR_BARE_RX.search(strip_comments(cached_text(p)))
        ]
        assert len(lane1) == 2, f"lane 1 裸形态期望 2 处，实得 {sorted(lane1)}"

    def test_the_two_faces_were_one_capability(
        self, k_files: list[pathlib.Path]
    ) -> None:
        """🔴 两个面调的是同一能力（同 `FormData` + 同 `file` 字段）⇒ 统一是对的。"""
        for p in k_files:
            if entry_of_path(p) not in LANE2:
                continue
            src = strip_comments(cached_text(p))
            if not OCR_WPID_RX.search(src):
                continue
            assert "FormData" in src, f"{p.name} 的 OCR 调用不是 multipart"
            assert re.search(r"append\(\s*'file'", src), (
                f"{p.name} 的 OCR 调用无 `file` 字段"
            )


class TestKBP40DerivedTotalInfixOnly:
    """derived_total 的 6 个「仅中置命中」键里本 lane 3 个；3 + 3 == 6。"""

    LANE2_INFIX = {
        "K11-2-total-occurrence",
        "K8-2-total-audited",
        "K9-2-total-audited",
    }

    def test_infix_only_set_is_six(self, k_files: list[pathlib.Path]) -> None:
        tail = derived_total_keys(k_files, include_infix=False)
        allk = derived_total_keys(k_files, include_infix=True)
        assert len(tail) == 77, f"尾部口径期望 77，实得 {len(tail)}"
        assert len(allk) == 83, f"全口径期望 83，实得 {len(allk)}"
        assert len(allk - tail) == 6, f"仅中置期望 6，实得 {sorted(allk - tail)}"

    def test_lane2_holds_exactly_three(self, k_files: list[pathlib.Path]) -> None:
        infix = derived_total_keys(k_files, include_infix=True) - derived_total_keys(
            k_files, include_infix=False
        )
        assert self.LANE2_INFIX <= infix, (
            f"缺 {sorted(self.LANE2_INFIX - infix)}"
        )
        got = {
            k
            for k in infix
            if re.match(r"^K(?:8|9|11|12|13)-", k)
        }
        assert got == self.LANE2_INFIX, f"本 lane 仅中置键 {sorted(got)}"

    def test_three_plus_three_equals_six(self, k_files: list[pathlib.Path]) -> None:
        """🔴 分摊对账：本 lane 3 + lane 1 的 3 == 6。"""
        infix = derived_total_keys(k_files, include_infix=True) - derived_total_keys(
            k_files, include_infix=False
        )
        l2 = {k for k in infix if re.match(r"^K(?:8|9|11|12|13)-", k)}
        l1 = {k for k in infix if re.match(r"^K(?:[1-7])-", k)}
        assert len(l2) == 3, sorted(l2)
        assert len(l1) == 3, sorted(l1)
        assert l2 | l1 == infix, f"有键落在两 lane 之外：{sorted(infix - l2 - l1)}"

    def test_infix_keys_are_middle_positioned_by_shape(self) -> None:
        """🔴 「仅中置」的形态：`total` 在键中段而非尾部。"""
        for k in self.LANE2_INFIX:
            assert "-total-" in k, f"{k} 不是中置形态"
            assert not k.endswith("-total"), f"{k} 是尾部形态"

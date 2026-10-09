# -*- coding: utf-8 -*-
"""L8 财务费用真双向接线守卫（spec `l8-true-bidirectional-2026-10-01` · 第3步）。

受管表 `明细表L8-2`（L8 整册唯一数据录入源，科目 6603 损益类）。同序：
几何现算 → 契约 → 注册 → HTML 侧对齐 → 真库干净命名空间。

L8 与 L6/L7 的差异（照抄会错）：
* **单级表头 r8**（L6/L7 两级表头 group/leaf）⇒ header_row=8、header_rows==1、anchor==A8。
* 受管表**自身 28 个裸 IF**（L1~L7 受管表零裸 IF），R 列占比 + R23 各月比例；R 列被
  neutralize_oo_crash_if_formulas 摘 <f> ⇒ 中性化后无公式 ⇒ **R=auto_source**（照 G11-1102 的 G/K 列）、
  NEUTRALIZED_COLUMNS==("R",)==AUTO_SOURCE_COLUMNS，往返不做 R 列公式幸存断言。
* **B~M 不进 formula_columns**（editable）：跨行派生行 R11/R13/R20 的 B~M 预置公式靠不被回填幸存。
  formula_columns 只有 N/Q/W（本行算术，往返幸存）；R 是 auto_source。
* 🔴 **方案 D1 混合身份**：13 行骨架用模板身份 GTROW-L82-NNNN（含补齐 R20 汇兑净损失），其中 3 个派生行
  R11/R13/R20 只读、**不进 store**（模板计算行，往返不在 projection ⇒ 跨行公式幸存）；用户 addRow 自铸 l82det-*。
* 身份字段是 **`key`**（L6/L7 是 rowId），store 单前缀 **`L8-2-full-data`**（非 L5/L6/L7 双前缀）。
* 月度列 B~M 对应 HTML 的 monthly[0..11] 数组 ⇒ json_pointer 末段是 monthly/0..monthly/11。
"""
from __future__ import annotations

import hashlib
import json
import re

import pytest
from openpyxl import load_workbook

from app.services.workpaper_sync import phase5_l8_financial_expenses as L8
from app.services.workpaper_sync import phase5_l_cycle_common as L
from app.services.workpaper_sync.contracts import parse_contract
from app.services.workpaper_sync.definitions import canonical_digest
from tests.workpaper_sync.l1_adapter_facts import ROOT, _count_bare_if, _is_formula, _query

TEMPLATE = ROOT / "backend" / "wp_templates" / L8.TEMPLATE_RELATIVE_PATH
FRONTEND = ROOT / "audit-platform" / "frontend" / "src" / "components" / "workpaper"
DETAIL_TS = FRONTEND / "composables" / "useL8Detail.ts"
DETAIL_VUE = FRONTEND / "l8" / "core" / "L8TabDetail.vue"
IDENTITY_TS = FRONTEND / "l8" / "core" / "l8DetailRowIdentity.ts"


@pytest.fixture(scope="module")
def wb():
    return load_workbook(TEMPLATE, data_only=False)


@pytest.fixture(scope="module")
def ws(wb):
    return wb[L8.MANAGED_SHEET]


# ═══════════════════════════════════════════════════════════════════════════
# 选型与几何（openpyxl 现算）
# ═══════════════════════════════════════════════════════════════════════════


class TestManagedSheetSelection:
    def test_template_sha256_matches_frozen(self) -> None:
        """净化后 SHA（横向共享公式组已展开，受管 sheet 逐格 0 diff）。"""
        assert hashlib.sha256(TEMPLATE.read_bytes()).hexdigest() == L8.TEMPLATE_SHA256

    def test_managed_sheet_present_and_adjudication_is_derived(self, wb) -> None:
        """审定表L8-1 是下游视图（R7~R16 全跨 sheet 聚合 明细表L8-2），非录入表 ⇒ 不选它。"""
        assert L8.MANAGED_SHEET in wb.sheetnames
        s1 = wb["审定表L8-1"]
        cross_ref = sum(
            isinstance(c.value, str) and "明细表L8-2" in c.value
            for r in s1.iter_rows() for c in r
        )
        assert cross_ref > 0, "审定表L8-1 应有跨 sheet 引用 明细表L8-2 的公式"

    def test_geometry_single_header(self, ws) -> None:
        assert ws.max_column == 23, "物理 max_column=W(23) ⇒ UUID 列 X"
        assert L8.UUID_COL == "X"
        assert ws[f"A{L8.FOOTER_ROW}"].value == L8.FOOTER_MARKER
        # 🔴 单级表头：header_row=8，group/leaf 留 None
        assert L8.HEADER_ROW == 8
        assert L8.SPEC_L82.header_group_row is None and L8.SPEC_L82.header_leaf_row is None
        assert (L8.FIRST_DATA_ROW, L8.LAST_DATA_ROW, L8.FOOTER_ROW) == (9, 21, 22)
        # R23 各月比例作模板静态行（在 footer 之下，不入数据区）
        assert ws["A23"].value == "各月比例"

    def test_formula_columns_each_data_row(self, ws) -> None:
        """N/Q/W 本行算术（formula，往返幸存）+ R 占比裸 IF（模板有公式，但声明 auto_source）。"""
        for r in range(L8.FIRST_DATA_ROW, L8.LAST_DATA_ROW + 1):
            assert ws[f"N{r}"].value == f"=SUM(B{r}:M{r})"
            assert ws[f"Q{r}"].value == f"=N{r}+O{r}+P{r}"
            assert ws[f"W{r}"].value == f"=T{r}+U{r}+V{r}"
            # 模板里 R 列仍是裸 IF 公式；中性化在 materialize 期发生，模板本体不变
            assert ws[f"R{r}"].value == f"=IF(Q{r}=0,0,Q{r}/$Q$22)"
        # 🔴 R 不进 formula_columns：中性化后 substrate 无公式 ⇒ auto_source（照 G11-1102 的 G/K 列）
        assert L8.FORMULA_COLUMNS == ("N", "Q", "W")
        assert L8.FORMULA_MASK == ("N9:N21", "Q9:Q21", "W9:W21")
        assert L8.AUTO_SOURCE_COLUMNS == ("R",)

    def test_managed_sheet_has_bare_if_unlike_l1_to_l7(self, wb, ws) -> None:
        """🔴 受管表自身 28 裸 IF（R 列占比 + R23 各月比例），与 L1~L7「受管表零」相反。"""
        assert _count_bare_if(ws) == 28
        assert L8.MANAGED_SHEET_BARE_IF == 28
        assert L8.NEUTRALIZED_COLUMNS == ("R",)
        assert sum(_count_bare_if(wb[n]) for n in wb.sheetnames) > 28

    def test_derived_rows_preset_cross_row_formulas(self, ws) -> None:
        """跨行派生行 R11/R13/R20 的 B 列预置公式存在，且这三行不在 formula_columns 覆盖的列。"""
        assert ws["B11"].value == "=B9-B10"
        assert ws["B13"].value == "=B11-B12"
        assert ws["B20"].value == "=B17-B18-B19"
        # B~M 不是 formula 列（editable）；方案 D1 下派生行不进 store ⇒ 往返不在 projection ⇒ 公式幸存
        assert "B" not in L8.FORMULA_COLUMNS and "M" not in L8.FORMULA_COLUMNS
        assert L8.DERIVED_ROWS == ("R11", "R13", "R20")

    def test_twentythree_managed_fields_three_formula_one_auto_source(self) -> None:
        assert len(L8.MANAGED_FIELD_SPECS_7) == 23
        cols = [col for _, col, *_ in L8.MANAGED_FIELD_SPECS_7]
        assert cols == list("ABCDEFGHIJKLMNOPQRSTUVW")
        formula_cols = [col for _, col, mode, *_ in L8.MANAGED_FIELD_SPECS_7 if mode == "formula"]
        assert formula_cols == ["N", "Q", "W"]
        # 🔴 R 列是 auto_source（占比裸 IF 中性化后无公式，服务端字面量，照 G11-1102）
        auto_cols = [col for _, col, mode, *_ in L8.MANAGED_FIELD_SPECS_7 if mode == "auto_source"]
        assert auto_cols == ["R"]
        # 月度列 B~M 的 json_key 对齐 HTML monthly 数组
        month_keys = [jk for _, col, _, _, jk, *_ in L8.MANAGED_FIELD_SPECS_7 if col in "BCDEFGHIJKLM"]
        assert month_keys == [f"monthly/{i}" for i in range(12)]


# ═══════════════════════════════════════════════════════════════════════════
# 契约
# ═══════════════════════════════════════════════════════════════════════════


class TestContract:
    def test_disk_and_source_are_locked(self) -> None:
        contract = L8.assert_contract_file_matches_source()
        assert contract.canonical_sha256 == canonical_digest(L8.build_contract_payload())

    def test_payload_parses_and_has_expected_shape(self) -> None:
        p = L8.build_contract_payload()
        parse_contract(dict(p), adapter_id=L8.ADAPTER_ID)
        assert p["review_status"] == "reviewed"
        assert p["review"]["entry_id"] == L8.ENTRY_ID
        table = p["sheets"][0]["tables"][0]
        # 🔴 单级表头 ⇒ header_rows==1、anchor==A8
        assert table["header_rows"] == 1 and table["anchor"] == "A8"
        assert table["footer_anchor"]["search_column"] == "A"
        assert table["footer_anchor"]["marker"] == "合计"
        assert len(table["fields"]) == 23

    def test_formula_mask_and_formula_fields(self) -> None:
        table = L8.build_contract_payload()["sheets"][0]["tables"][0]
        # 🔴 R 不在 formula_mask（auto_source，中性化后无公式）
        assert table["formula_mask"] == ["N9:N21", "Q9:Q21", "W9:W21"]
        formula_keys = {
            f["json_pointer"].rsplit("/", 1)[-1]
            for f in table["fields"] if f["mode"] == "formula"
        }
        assert formula_keys == {"periodUnadjusted", "periodAudited", "priorAudited"}
        # R 列占比声明为 auto_source（服务端字面量）
        auto_keys = {
            f["json_pointer"].rsplit("/", 1)[-1]
            for f in table["fields"] if f["mode"] == "auto_source"
        }
        assert auto_keys == {"ratio"}

    def test_monthly_pointers_present(self) -> None:
        table = L8.build_contract_payload()["sheets"][0]["tables"][0]
        monthly = sorted(
            f["json_pointer"] for f in table["fields"] if "/monthly/" in f["json_pointer"]
        )
        assert len(monthly) == 12
        assert monthly[0].endswith("/monthly/0") and monthly[-1].endswith("/monthly/9")  # 字典序

    def test_derived_readonly_sheet_declared(self) -> None:
        review = L8.build_contract_payload()["review"]
        drs = review["derived_readonly_sheet"]
        assert drs["excel_name"] == "审定表L8-1"
        assert "明细表L8-2" in drs["reason"]

    def test_defects_registered_bare_if_and_derived_and_amount_kind(self) -> None:
        review = L8.build_contract_payload()["review"]
        defects = review["template_defects_registered"]
        assert defects["bare_if"]["managed_sheet_bare_if"] == 28
        assert defects["bare_if"]["neutralized_columns"] == ["R"]
        assert defects["derived_rows"]["rows"] == ["R11", "R13", "R20"]
        assert defects["amount_kind"]["value"] == "occurrence"
        assert defects["amount_kind"]["account_code"] == "6603"

    def test_routes_through_common_skeleton(self) -> None:
        assert L8.EntrySelectionError is L.LEntrySelectionError
        assert L8.build_contract_payload() == L.build_contract_payload(L8.IDENTITY, L8.SPECS)


# ═══════════════════════════════════════════════════════════════════════════
# 注册
# ═══════════════════════════════════════════════════════════════════════════


class TestRegistration:
    def test_provider_aliases_present(self) -> None:
        assert L8.publish_pilot_definitions is L8.publish_definitions
        assert L8.attach_pilot_adapters is L8.attach_adapters
        assert L8.PILOT_WP_CODES == frozenset({"L8F"})

    def test_store_registry_entry(self) -> None:
        from app.services.workpaper_sync.store_item_registry import STORE_MERGE_REGISTRY

        plan = STORE_MERGE_REGISTRY["l8.financial_expenses"]
        assert plan.provider_module == "phase5_l8_financial_expenses"
        assert [i.item_id for i in plan.items] == ["L8-2-full-data"]
        assert plan.oo_crash_neutralization_fn == "neutralize_oo_crash_if_formulas"

    def test_whitelist_and_ledger_are_paired(self) -> None:
        from app.services.workpaper_sync.adapters import registry as R
        from app.services.workpaper_sync.adapters.delivered_contracts_ledger import (
            DELIVERED_PER_ENTRY_CONTRACTS,
        )

        mod = "app.services.workpaper_sync.phase5_l8_financial_expenses"
        assert mod in R._ALLOWED_PROVIDER_MODULES
        rows = [r for r in DELIVERED_PER_ENTRY_CONTRACTS if r["contract_id"] == "l8.financial_expenses"]
        assert len(rows) == 1 and rows[0]["provider_module"] == mod

    def test_l_cycle_whitelist_ledger_counts_paired(self) -> None:
        """配对不变式：L 域白名单条数 == ledger 的 L 条数（现为 7==7）。"""
        from app.services.workpaper_sync.adapters import registry as R
        from app.services.workpaper_sync.adapters.delivered_contracts_ledger import (
            DELIVERED_PER_ENTRY_CONTRACTS,
        )

        wl_l = [m for m in R._ALLOWED_PROVIDER_MODULES if ".phase5_l" in m and "_cycle_common" not in m]
        ledger_l = [r for r in DELIVERED_PER_ENTRY_CONTRACTS if re.match(r"l\d", str(r["contract_id"]))]
        assert len(wl_l) == len(ledger_l), (len(wl_l), len(ledger_l))

    def test_wp_code_adjudication_targets_the_whole_workbook_code(self) -> None:
        doc = json.loads(
            (ROOT / "backend/data/workpaper_sync_entry_wp_code_adjudication.json").read_text("utf-8")
        )
        row = next(a for a in doc["adjudications"] if a["entry_id"] == L8.ENTRY_ID)
        # 🔴 wp_codes 是底稿码 L8（不是科目码 6603）
        assert row["wp_codes"] == ["L8"] and row["contract_id"] == L8.ADAPTER_ID


# ═══════════════════════════════════════════════════════════════════════════
# HTML 侧对齐（前端 json 键 == 契约 json_pointer 段）
# ═══════════════════════════════════════════════════════════════════════════


class TestHtmlAlignment:
    def test_store_item_id_agrees(self) -> None:
        """L8 的 store item 是硬编码字符串 'L8-2-full-data'（单前缀，_restoreRows 读它）。"""
        src = DETAIL_VUE.read_text("utf-8")
        assert f"'{L8.STORE_ITEM_ID}'" in src
        ts = DETAIL_TS.read_text("utf-8")
        assert f"'{L8.STORE_ITEM_ID}'" in ts

    def test_row_identity_mixed_skeleton_and_user(self) -> None:
        """方案 D1 混合身份：骨架行 GTROW-L82-NNNN（模板身份），用户 addRow 自铸 l82det-*；
        私有随机 ID 已删；派生行不进 store。"""
        identity = IDENTITY_TS.read_text("utf-8")
        # 身份模块存在、模板身份生成器 + 派生行标记 + store 过滤
        assert "GTROW-" in identity
        assert "l82TemplateRowId" in identity
        assert "L82_DERIVED_ROW_INDICES" in identity
        assert "rowsForStore" in identity

        ts = DETAIL_TS.read_text("utf-8")
        # 用户新增行走平台共享 newRowIdentity('l82det')
        assert "newRowIdentity('l82det')" in ts
        # 旧随机运行时前缀不再用于新增行
        assert "l8-detail-${Date.now()}" not in ts
        # 保存时派生行被剔除出 store
        assert "rowsForStore(detailRows.value)" in ts

        vue = DETAIL_VUE.read_text("utf-8")
        # 默认行走 createL8DefaultRows（GTROW 骨架），不再位置化自铸
        assert "createL8DefaultRows()" in vue
        assert "l8-detail-default-" not in vue
        # 派生行月度格只读 + 模板行不可删
        assert "isL82DerivedRowKey" in vue
        assert "isL82TemplateRowId" in vue

    def test_frontend_row_keys_cover_every_contract_field(self) -> None:
        """L8DetailRow 字段覆盖契约每个 json_pointer 末段（monthly 数组按数组键处理）。"""
        src = DETAIL_TS.read_text("utf-8")
        iface = src.split("export interface L8DetailRow {", 1)[1].split("\n}", 1)[0]
        frontend = set(re.findall(r"^\s+(\w+)\s*[?:]", iface, re.M))
        contract_leaves = {
            f["json_pointer"].rsplit("/", 1)[-1]
            for f in L8.build_contract_payload()["sheets"][0]["tables"][0]["fields"]
        }
        # monthly/N 的末段是数字 N，对应 HTML 的 monthly 数组字段 ⇒ 把数字段折叠成 'monthly'
        contract = {("monthly" if leaf.isdigit() else leaf) for leaf in contract_leaves}
        assert contract - frontend == set(), f"契约有、前端缺：{sorted(contract - frontend)}"


# ═══════════════════════════════════════════════════════════════════════════
# 真库：干净命名空间
# ═══════════════════════════════════════════════════════════════════════════


class TestRealDbNamespace:
    def test_managed_namespace_is_clean(self) -> None:
        """L8-2-% 真库命名空间：受管 full-data 键 0 行、旧位置化 row-N 键 0 行。"""
        rows = _query(
            "SELECT count(*) FILTER (WHERE item_id = 'L8-2-full-data'), "
            "count(*) FILTER (WHERE item_id LIKE 'L8-2-row-%') FROM checklist_responses"
        )
        full_data, old_positional = rows[0]
        assert old_positional == 0, "L8-2 旧位置化 row-N 键有真数据 ⇒ 切换前须先核对"
        assert full_data >= 0

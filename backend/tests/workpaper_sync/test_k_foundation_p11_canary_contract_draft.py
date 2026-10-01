# -*- coding: utf-8 -*-
"""K 循环 foundation spec — 阶段 2 Task 22：K10 per-entry contract 草案。

spec: k-cycle-sync-foundation-and-first-canary
Task 22: 发 K10 per-entry contract（`[ ]*` 依赖 BP-1）
Property: KF-P48, KF-P24, KF-P37

═══ 🔴 `[ ]*` 的边界 ═══

spec 明文：「发布生产契约需 approved authority model（BP-1），本任务只交付
**契约草案 + 判据**」。

⇒ 本文件交付的 `k10.other_income_adjustment.candidate.json` 标
`review_status: "candidate"`，**不是**生产契约。判据两侧都验：
  - 草案的结构与已 reviewed 契约同构（schema 可用）
  - 草案**不得**被当生产契约注册（review_status != reviewed）

═══ 契约里落地的三个 spec 裁决 ═══

1. **KF-P48**：显式排除 6 个 `K10-review-session-*` 键（另一命名空间）
   🔴 **本轮新增第四类**：`K10-3-adjustment` 是 AI section-id（design.md
   的 KC-5 三命名空间表未登记）
2. **KF-P24**：`row_delete_api_kind` 用四元组（不是单值枚举）
3. **KF-P37**：derived_total 键声明（K10 侧属 83 个全集的子集）

═══ 🔴 承接 Task 20 的发现 ═══

契约的 `footer_anchor.carries_total_formula = false` + `why_no_footer` 显式
说明「13 册「调整分录汇总」SUM 数全 0」⇒ 金额判据按行级不按 footer。
"""
from __future__ import annotations

import hashlib
import json
import pathlib
import re

import pytest
from openpyxl import load_workbook

from tests.workpaper_sync.k_foundation_facts import (  # noqa: E402
    DATA,
    ROOT,
    cached_text,
    derived_total_keys,
    k_domain_files,
    strip_comments,
)

CONTRACT_DIR = DATA / "workpaper_sync_contracts"
#: 🔴 2026-10-01 草案已被 reviewed 生产契约取代（K8/K9/K11/K12/K13 已真双向发布）。
#:    草案本体作为 Task 18/22「草案阶段」的交付证据归档在 spec evidence 目录，本文件的
#:    草案判据改读归档（历史事实不变）；生产契约的判据见 `test_k_cycle_reviewed_contracts.py`。
DRAFT_ARCHIVE_DIR = (
    ROOT / ".kiro" / "specs" / "k8-k9-k11-k12-k13-dedicated-composable-and-cross-cycle-hub"
    / "evidence" / "superseded-candidate-contracts"
)
DRAFT_PATH = DRAFT_ARCHIVE_DIR / "k10.other_income_adjustment.candidate.json"
K_TEMPLATE_DIR = ROOT / "backend" / "wp_templates" / "K"
CANARY_WORKBOOK = "K10 其他收益.xlsx"
CANARY_SHEET = "调整分录汇总K10-3"
CANARY_KEY = "K10-3-entries"

#: 模板 R5 表头十列（现算）
TEMPLATE_HEADERS = {
    "A": "调整事项说明",
    "B": "类别（报表调整/账项调整/其他）",
    "C": "报表项目",
    "D": "科目名称",
    "E": "附注项目",
    "F": "……",
    "G": "借方调整金额",
    "H": "贷方调整金额",
    "I": "索引",
    "J": "备注",
}

#: 前端 `AdjustmentEntry` 的 11 个字段
FRONTEND_FIELDS = (
    "id", "type", "description", "category", "reportItem",
    "accountName", "noteItem", "debitAmount", "creditAmount",
    "refIndex", "remark",
)


@pytest.fixture(scope="module")
def draft() -> dict:
    assert DRAFT_PATH.exists(), f"契约草案不存在：{DRAFT_PATH}"
    return json.loads(DRAFT_PATH.read_text(encoding="utf-8"))


@pytest.fixture(scope="module")
def k_files() -> list[pathlib.Path]:
    return k_domain_files()


@pytest.fixture(scope="module")
def managed_table(draft: dict) -> dict:
    return draft["sheets"][0]["tables"][0]


# ════════════════════════════════════════════════════════════════════════════
# `[ ]*` 边界：草案不是生产契约
# ════════════════════════════════════════════════════════════════════════════
class TestDraftIsNotAProductionContract:
    """🔴 BP-1 未交付 ⇒ 草案标 candidate，不得被当生产契约。"""

    def test_review_status_is_candidate(self, draft: dict) -> None:
        assert draft["review_status"] == "candidate", (
            f"草案的 review_status 是 {draft['review_status']!r}"
            " ⇒ 它会被当生产契约注册，违反 `[ ]*` 边界"
        )

    def test_semantic_version_marks_draft(self, draft: dict) -> None:
        assert "draft" in draft["semantic_version"], (
            f"semantic_version {draft['semantic_version']!r} 未标 draft"
        )

    def test_why_candidate_is_explicit(self, draft: dict) -> None:
        """显式登记为什么是 candidate（不是忘了改）。"""
        why = draft["review"].get("why_candidate_not_reviewed", "")
        assert why.strip(), "缺 why_candidate_not_reviewed"
        assert "BP-1" in why, "未指明卡 BP-1"

    def test_production_contracts_are_exactly_the_promotion_ledger(self) -> None:
        """🔴 晋级后：K 循环 reviewed 生产契约 **恰等于** 晋级账本（多一条少一条都红）。

        原判据是「K 循环 0 生产契约」（晋级前的空分母登记）。2026-10-01 六册调整分录汇总
        发了 reviewed 契约 ⇒ 判据翻面为「现算 == `K_REVIEWED_CONTRACTS`」，而不是删掉 ——
        删掉就无从发现「有人顺手又给 K 发了一份未登记的契约」。
        """
        from tests.workpaper_sync.k_foundation_facts import (
            K_REVIEWED_CONTRACTS,
            k_reviewed_contract_owners,
        )

        live = k_reviewed_contract_owners()
        expected = {f"{cid}.json" for cid in K_REVIEWED_CONTRACTS.values()}
        assert set(live) == expected, (
            f"K reviewed 契约现算 {sorted(live)} ≠ 账本 {sorted(expected)}"
        )
        for n, cid in K_REVIEWED_CONTRACTS.items():
            assert live[f"{cid}.json"].startswith(f"xlsx/gt-k{n}-"), (
                f"{cid} 归属 {live[f'{cid}.json']} 与序号 K{n} 不符"
            )

    def test_draft_is_superseded_not_left_beside_the_reviewed_contract(self) -> None:
        """🔴 草案已归档，契约目录里**不得**再有同 adapter 的 candidate（双源会让门失效）。"""
        assert DRAFT_PATH.exists(), "草案归档缺失 ⇒ 草案阶段交付证据丢了"
        live_candidate = CONTRACT_DIR / DRAFT_PATH.name
        assert not live_candidate.exists(), (
            f"{live_candidate.name} 仍在契约目录 ⇒ 与 reviewed 契约双源"
        )
        assert (CONTRACT_DIR / "k10.other_income_adjustment.json").exists()

    def test_entry_id_must_be_null_for_candidate(self, draft: dict) -> None:
        """🔴 candidate 的 `review.entry_id` **必须为 null**。

        `test_task53_k_cycle_migration.py` 的两条判据
        （`test_no_slice_entry_has_a_contract` /
        `test_no_pilot_contract_belongs_to_the_k_cycle`）按 `review.entry_id`
        判归属且**不区分** candidate/reviewed ⇒ 草案若填 entry_id 会打破它们。
        目标 entry 记在 `draft_target_entry_id`。
        """
        assert draft["review"]["entry_id"] is None, (
            f"candidate 契约的 entry_id 是 {draft['review']['entry_id']!r}"
            " ⇒ 会打破 test_task53 的两条既有判据"
        )
        assert draft["review"]["draft_target_entry_id"] == (
            "xlsx/gt-k10-other-income"
        )
        assert draft["review"]["why_entry_id_is_null"].strip()

    def test_this_matches_the_example_candidate_caliber(self) -> None:
        """两侧都验：与 `_example.candidate.json` 同口径。"""
        example = json.loads(
            (CONTRACT_DIR / "_example.candidate.json").read_text(encoding="utf-8")
        )
        assert example["review_status"] == "candidate"
        assert (example.get("review") or {}).get("entry_id") is None, (
            "参考 candidate 的 entry_id 非 null ⇒ 口径依据变了"
        )

    def test_only_ledger_contracts_claim_k_entries(self) -> None:
        """🔴 契约目录里声明 K entry 归属的，恰是晋级账本那 6 份（不论 status）。"""
        from tests.workpaper_sync.k_foundation_facts import K_REVIEWED_CONTRACTS

        slice_doc = json.loads(
            (DATA / "workpaper_sync_k_cycle_manifest_slice.json").read_text(
                encoding="utf-8"
            )
        )
        k_ids = {e["entry_id"] for e in slice_doc["independent_entries"]}
        owners: set[str] = set()
        for p in sorted(CONTRACT_DIR.glob("*.json")):
            doc = json.loads(p.read_text(encoding="utf-8"))
            owner = (doc.get("review") or {}).get("entry_id")
            if owner in k_ids:
                owners.add(p.name)
        assert owners == {f"{cid}.json" for cid in K_REVIEWED_CONTRACTS.values()}, owners


# ════════════════════════════════════════════════════════════════════════════
# 草案与已 reviewed 契约同构（schema 可用）
# ════════════════════════════════════════════════════════════════════════════
class TestDraftSchemaIsIsomorphicToReviewed:
    """草案的顶层结构与 h9 契约同构 ⇒ schema 可用不是自创。"""

    REFERENCE = "h9.lease_liability_detail.json"

    def test_top_level_keys_are_a_subset_of_reference(self, draft: dict) -> None:
        ref = json.loads(
            (CONTRACT_DIR / self.REFERENCE).read_text(encoding="utf-8")
        )
        ref_keys = set(ref)
        draft_keys = set(draft)
        extra = draft_keys - ref_keys
        assert extra == set(), (
            f"草案有参考契约没有的顶层键：{sorted(extra)} ⇒ schema 自创"
        )

    def test_schema_version_matches_reference(self, draft: dict) -> None:
        ref = json.loads(
            (CONTRACT_DIR / self.REFERENCE).read_text(encoding="utf-8")
        )
        assert draft["schema_version"] == ref["schema_version"]

    #: 🔴 KC-7 要求**新增**的 table 级字段（参考契约还没有）
    KC7_NEW_TABLE_KEYS = {"row_delete_api_kind"}

    def test_table_level_keys_are_a_subset_plus_kc7_addition(
        self, managed_table: dict
    ) -> None:
        """table 级键 ⊆ 参考契约 ∪ KC-7 新增字段。

        🔴 `row_delete_api_kind` 是 KC-7 明文要求新增的四元组字段
        （「单值枚举装不下 13 种签名，SHALL 改为四元组」）。参考契约 h9
        还没有它 ⇒ 差集恰好是这一个键，不是 schema 自创。
        """
        ref = json.loads(
            (CONTRACT_DIR / self.REFERENCE).read_text(encoding="utf-8")
        )
        ref_table = ref["sheets"][0]["tables"][0]
        extra = set(managed_table) - set(ref_table)
        assert extra == self.KC7_NEW_TABLE_KEYS, (
            f"差集期望恰好 {sorted(self.KC7_NEW_TABLE_KEYS)}，实得 {sorted(extra)}"
            " ⇒ 除 KC-7 新增字段外有自创键"
        )

    def test_kc7_new_field_is_absent_from_all_reviewed_contracts(
        self,
    ) -> None:
        """🔴 两侧都验：没有任何 reviewed 契约已有 `row_delete_api_kind`。

        这证明它确实是**新增**字段（KC-7 的裁决尚未在平台落地），
        不是我漏抄了参考契约的既有字段。
        """
        has_it: list[str] = []
        for p in sorted(CONTRACT_DIR.glob("*.json")):
            doc = json.loads(p.read_text(encoding="utf-8"))
            if doc.get("review_status") != "reviewed":
                continue
            for sheet in doc.get("sheets", []):
                for table in sheet.get("tables", []):
                    if "row_delete_api_kind" in table:
                        has_it.append(p.name)
        assert has_it == [], (
            f"已有 reviewed 契约带 row_delete_api_kind：{has_it}"
            " ⇒ 应改用它们的口径而不是新造"
        )

    def test_field_level_keys_are_a_subset(self, managed_table: dict) -> None:
        ref = json.loads(
            (CONTRACT_DIR / self.REFERENCE).read_text(encoding="utf-8")
        )
        ref_field_keys = set(ref["sheets"][0]["tables"][0]["fields"][0])
        for f in managed_table["fields"]:
            extra = set(f) - ref_field_keys
            assert extra == set(), (
                f"字段 {f['column_key']} 有参考没有的键：{sorted(extra)}"
            )


# ════════════════════════════════════════════════════════════════════════════
# KF-P48：managed table 与三个排除命名空间
# ════════════════════════════════════════════════════════════════════════════
class TestKFP48ManagedTableAndExclusions:
    """🔴 显式排除 review-session 键 + 本轮新增的 AI section-id 类。"""

    def test_managed_table_is_the_canary_key(self, draft: dict) -> None:
        assert draft["review"]["html_store"]["item_ids"] == [CANARY_KEY]

    def test_review_session_keys_are_excluded(self, draft: dict) -> None:
        excl = draft["review"]["excluded_namespaces"]
        rs = excl["review_session"]
        assert rs["live_count"] == 6, (
            f"review-session 键数期望 6，草案写 {rs['live_count']}"
        )
        assert "14位时间戳" in rs["keys_pattern"]
        assert "managed table" in rs["why"]

    def test_publish_flag_is_excluded(self, draft: dict) -> None:
        excl = draft["review"]["excluded_namespaces"]
        assert excl["publish_flag"]["keys"] == ["K10-3-published"]

    def test_ai_section_id_is_excluded_as_the_fourth_namespace(
        self, draft: dict
    ) -> None:
        """🔴 本轮新增：`K10-3-adjustment` 是 AI section-id（KC-5 表未登记）。"""
        excl = draft["review"]["excluded_namespaces"]
        assert "ai_section_id" in excl, (
            "草案没排除 AI section-id 类 ⇒ 第四类命名空间会被当业务键"
        )
        assert excl["ai_section_id"]["keys"] == ["K10-3-adjustment"]
        why = excl["ai_section_id"]["why"]
        assert "KC-5" in why, "未指明它是 KC-5 表的缺口"
        assert "checklist" in why, "未说明它不在持久化路径"

    def test_all_three_k10_3_keys_are_accounted_for(
        self, draft: dict, k_files: list[pathlib.Path]
    ) -> None:
        """K10-3 组的 3 个键：1 个 managed + 2 个排除，无遗漏。"""
        from tests.workpaper_sync.k_foundation_facts import business_keys
        k3_keys = {
            k for k in business_keys(k_files)
            if k.startswith("K10-3-")
        }
        managed = set(draft["review"]["html_store"]["item_ids"])
        excl = draft["review"]["excluded_namespaces"]
        excluded = set(excl["publish_flag"]["keys"]) | set(
            excl["ai_section_id"]["keys"]
        )
        assert managed | excluded == k3_keys, (
            f"K10-3 组有未归类的键：{sorted(k3_keys - managed - excluded)}"
        )
        assert managed & excluded == set()


# ════════════════════════════════════════════════════════════════════════════
# KF-P24：row_delete_api_kind 四元组
# ════════════════════════════════════════════════════════════════════════════
class TestKFP24RowDeleteQuadrupleInContract:
    """🔴 KC-7：契约字段须是四元组，不是单值枚举。"""

    def test_row_delete_api_kind_is_a_quadruple(
        self, managed_table: dict
    ) -> None:
        rd = managed_table["row_delete_api_kind"]
        assert isinstance(rd, dict), (
            f"row_delete_api_kind 是 {type(rd).__name__} ⇒ 单值枚举装不下 13 种签名"
        )
        for key in ("kind", "arity", "param_order", "param_name_family"):
            assert key in rd, f"四元组缺 {key}"

    def test_arity_and_param_order_are_consistent(
        self, managed_table: dict
    ) -> None:
        rd = managed_table["row_delete_api_kind"]
        assert rd["arity"] == len(rd["param_order"]), (
            f"arity {rd['arity']} 与 param_order {rd['param_order']} 不一致"
        )

    def test_kind_is_by_identity_not_by_index(
        self, managed_table: dict
    ) -> None:
        """🔴 K10 按身份删（不是按下标）—— 它不在下标族 K2~K7 里。"""
        rd = managed_table["row_delete_api_kind"]
        assert rd["kind"] == "by_identity", (
            f"kind 是 {rd['kind']!r} ⇒ K10 若按下标删会踩 BP-8 的组合风险"
        )


# ════════════════════════════════════════════════════════════════════════════
# KF-P37：derived_total 声明
# ════════════════════════════════════════════════════════════════════════════
class TestKFP37DerivedTotalInContract:
    """契约声明 derived_total 键（K10 侧属 83 个全集的子集）。"""

    def test_k10_derived_total_keys_are_a_subset_of_83(
        self, k_files: list[pathlib.Path]
    ) -> None:
        full = derived_total_keys(k_files, include_infix=True)
        assert len(full) == 83
        k10_derived = {k for k in full if re.match(r"^K10(?![0-9])-", k)}
        assert k10_derived <= full
        # K10 有 derived_total 键（非空分母）
        assert k10_derived, (
            "K10 侧 derived_total 为空 ⇒ 契约的 derived 声明无对象"
        )

    def test_canary_key_itself_is_not_a_derived_total(
        self, k_files: list[pathlib.Path]
    ) -> None:
        """canary 是行数据不是派生合计。"""
        full = derived_total_keys(k_files, include_infix=True)
        assert CANARY_KEY not in full, (
            f"{CANARY_KEY} 被判为 derived_total ⇒ 它不该作 managed table"
        )

    def test_contract_declares_no_derived_fields_for_canary(
        self, draft: dict
    ) -> None:
        """🔴 canary 的数据区无公式 ⇒ `derived_fields` 为空、`formula_mask` 为空。"""
        assert draft["review"]["derived_fields"] == [], (
            "canary 有 derived_fields ⇒ 与「数据区纯空白无公式」矛盾"
        )

    def test_formula_mask_is_empty(self, managed_table: dict) -> None:
        assert managed_table["formula_mask"] == [], (
            "canary 有 formula_mask ⇒ 与「数据区无公式」矛盾"
        )


# ════════════════════════════════════════════════════════════════════════════
# 🔴 承接 Task 20：footer_anchor 声明无 SUM
# ════════════════════════════════════════════════════════════════════════════
class TestFooterAnchorDeclaresNoSum:
    """🔴 契约必须显式声明「本表无 footer SUM」，否则 merge 会找不到合计行。"""

    def test_carries_total_formula_is_false(self, managed_table: dict) -> None:
        fa = managed_table["footer_anchor"]
        assert fa["carries_total_formula"] is False, (
            "契约声称有合计公式 ⇒ 与实测（13 册 SUM 全 0）矛盾"
        )

    def test_why_no_footer_is_documented(self, managed_table: dict) -> None:
        why = managed_table["footer_anchor"].get("why_no_footer", "")
        assert why.strip(), "缺 why_no_footer 说明"
        assert "SUM" in why, "未说明 SUM 为 0"
        assert "行级" in why, "未指出金额判据按行级"

    def test_template_really_has_no_sum(self) -> None:
        """两侧都验：模板确实无 SUM。"""
        wb = load_workbook(
            K_TEMPLATE_DIR / CANARY_WORKBOOK, read_only=False, data_only=False
        )
        try:
            ws = wb[CANARY_SHEET]
            sums = [
                c.coordinate for row in ws.iter_rows() for c in row
                if isinstance(c.value, str) and "SUM(" in c.value.upper()
            ]
            assert sums == [], f"模板有 SUM {sums} ⇒ 契约声明须改"
        finally:
            wb.close()

    def test_balance_check_is_in_application_layer(
        self, k_files: list[pathlib.Path]
    ) -> None:
        """借贷合计由前端 `currentBalance` computed 校验（不依赖模板公式）。"""
        tab = next(
            (p for p in k_files if p.name == "K10TabAdjustment.vue"), None
        )
        assert tab is not None
        src = strip_comments(cached_text(tab))
        assert "currentBalance" in src, (
            "找不到 currentBalance ⇒ 应用层借贷校验的登记失效"
        )
        assert re.search(r"totalDebit\s*-\s*totalCredit", src), (
            "找不到 totalDebit - totalCredit ⇒ 借贷平衡判据不在应用层"
        )


# ════════════════════════════════════════════════════════════════════════════
# 字段映射三边一致（契约 ↔ 模板 ↔ 前端）
# ════════════════════════════════════════════════════════════════════════════
class TestFieldMappingIsTriadConsistent:
    """契约字段 ↔ 模板列头 ↔ 前端字段名 三边逐一对应。"""

    def test_nine_mapped_fields(self, managed_table: dict) -> None:
        """9 个映射列（10 列减去占位列 F）。"""
        fields = managed_table["fields"]
        assert len(fields) == 9, (
            f"映射字段数期望 9，实得 {len(fields)}"
        )
        cols = {f["cell"]["column"] for f in fields}
        assert cols == set("ABCDEGHIJ"), (
            f"映射列集合不符：{sorted(cols)}（F 是占位列须排除）"
        )

    def test_header_text_matches_template_byte_for_byte(
        self, managed_table: dict
    ) -> None:
        """🔴 表头逐字节一致（含全角括号与省略号）。"""
        wb = load_workbook(
            K_TEMPLATE_DIR / CANARY_WORKBOOK, read_only=False, data_only=False
        )
        try:
            ws = wb[CANARY_SHEET]
            for f in managed_table["fields"]:
                col = f["cell"]["column"]
                actual = ws[f"{col}5"].value
                assert f["header_text"] == actual, (
                    f"{col}5 表头不符：契约 {f['header_text']!r} vs 模板 {actual!r}"
                )
                assert TEMPLATE_HEADERS[col] == actual, (
                    f"{col}5 与现算基线不符"
                )
        finally:
            wb.close()

    def test_json_pointers_match_frontend_field_names(
        self, managed_table: dict, k_files: list[pathlib.Path]
    ) -> None:
        """契约的 json_pointer 末段 == 前端 `AdjustmentEntry` 字段名。"""
        tab = next(p for p in k_files if p.name == "K10TabAdjustment.vue")
        src = strip_comments(cached_text(tab))
        for f in managed_table["fields"]:
            field_name = f["json_pointer"].rsplit("/", 1)[-1]
            assert re.search(rf"\b{re.escape(field_name)}\b", src), (
                f"前端找不到字段 {field_name}（来自 {f['json_pointer']}）"
            )

    def test_type_is_store_only(self, draft: dict) -> None:
        """`type`(AJE/RJE) 是 store-only（模板无列）。"""
        assert draft["review"]["store_only_fields"] == ["type"]
        # 模板表头里确实没有「类型」列
        assert "type" not in TEMPLATE_HEADERS.values()

    def test_placeholder_column_f_is_template_only(self, draft: dict) -> None:
        """模板 F 列（表头 `……`）是 template-only。"""
        tol = draft["review"]["template_only_columns"]
        assert len(tol) == 1
        assert tol[0]["column"] == "F"
        assert tol[0]["header_text"] == "……"

    def test_field_count_arithmetic(self, draft: dict, managed_table: dict) -> None:
        """算术自检：9 映射 + 1 store-only(type) + 1 身份(id) == 11 前端字段。"""
        mapped = len(managed_table["fields"])
        store_only = len(draft["review"]["store_only_fields"])
        identity = 1  # id
        assert mapped + store_only + identity == len(FRONTEND_FIELDS), (
            f"{mapped} + {store_only} + {identity} != {len(FRONTEND_FIELDS)}"
        )

    def test_amount_fields_are_typed_amount(self, managed_table: dict) -> None:
        """金额列 G/H 的 value_type 是 amount。"""
        by_col = {f["cell"]["column"]: f for f in managed_table["fields"]}
        for col in ("G", "H"):
            assert by_col[col]["value_type"] == "amount", (
                f"{col} 列 value_type 是 {by_col[col]['value_type']!r} 而非 amount"
            )

    def test_all_data_fields_are_editable(self, managed_table: dict) -> None:
        """🔴 数据区无公式 ⇒ 9 个字段全 editable（无 formula mode）。"""
        modes = {f["mode"] for f in managed_table["fields"]}
        assert modes == {"editable"}, (
            f"字段 mode 集合 {sorted(modes)} ⇒ 与「数据区无公式」矛盾"
        )


# ════════════════════════════════════════════════════════════════════════════
# 行身份声明
# ════════════════════════════════════════════════════════════════════════════
class TestRowIdentityDeclaration:
    """契约声明行身份形态与族。"""

    def test_row_identity_is_field_kind(self, managed_table: dict) -> None:
        ri = managed_table["row_identity"]
        assert ri["kind"] == "field"
        assert ri["json_pointer"] == "/rows/*/id"

    def test_pattern_and_family_are_declared(self, managed_table: dict) -> None:
        ri = managed_table["row_identity"]
        assert ri["pattern"] == "entry-{13位时间戳}-{6位base36}"
        assert ri["family"] == "c", (
            f"family 是 {ri['family']!r} ⇒ 与「安全族」矛盾"
        )
        assert "ENTROPY" in ri["why_safe"]
        assert "FALLBACK" in ri["why_safe"]

    def test_uuid_column_is_beyond_effective_columns(
        self, draft: dict, managed_table: dict
    ) -> None:
        """UUID 列在有效列之后（10 列 ⇒ UUID 在 K）。"""
        assert draft["review"]["effective_columns"] == 10
        assert managed_table["uuid_col"] == "K"
        assert draft["review"]["uuid_column"] == "K"


# ════════════════════════════════════════════════════════════════════════════
# 模板指纹与 second_writer
# ════════════════════════════════════════════════════════════════════════════
class TestTemplateFingerprintAndSecondWriter:

    def test_template_sha256_matches_disk(self, draft: dict) -> None:
        p = K_TEMPLATE_DIR / CANARY_WORKBOOK
        actual = hashlib.sha256(p.read_bytes()).hexdigest()
        assert draft["template"]["template_sha256"] == actual, (
            f"模板指纹漂移：契约 {draft['template']['template_sha256'][:16]}... "
            f"vs 磁盘 {actual[:16]}..."
        )

    def test_relative_path_is_correct(self, draft: dict) -> None:
        rel = draft["template"]["relative_path"]
        assert (ROOT / "backend" / "wp_templates" / rel).exists()

    def test_second_writer_is_declared(self, draft: dict) -> None:
        """🔴 `useAdjustmentCentralSync` 是第二写入方，契约须声明。"""
        sw = draft["review"]["second_writer"]
        assert sw["symbol"] == "useAdjustmentCentralSync"
        assert sw["sites"] == 3
        assert "写入方集合" in sw["why"]

    def test_second_writer_count_matches_source(
        self, draft: dict, k_files: list[pathlib.Path]
    ) -> None:
        """两侧都验：契约声明的 3 处与源码现算一致。"""
        tab = next(p for p in k_files if p.name == "K10TabAdjustment.vue")
        src = strip_comments(cached_text(tab))
        actual = len(re.findall(r"useAdjustmentCentralSync", src))
        assert draft["review"]["second_writer"]["sites"] == actual

    def test_tb_publish_gate_carrier_is_declared(self, draft: dict) -> None:
        """TB 发布门载体是 `useK10FormData.ts`（四层里的 FormData 层）。"""
        assert draft["review"]["tb_publish_gate"] == "useK10FormData.ts"
        assert draft["review"]["carrier"]["tb_publish_gate"] == "useK10FormData.ts"

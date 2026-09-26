"""判据：D1-5 `single_html` 裁决的四条依据现在仍然成立（防证据腐烂）。

spec: d1-sync-row-table-engine-and-d1-coverage · Task 27（D1-5 子目标）· Requirements 5.9 / 5.10
裁决: .kiro/specs/d1-sync-row-table-engine-and-d1-coverage/evidence/
      task27-d1-5-single-html-adjudication.json

裁决类证据最容易腐烂：一旦有人给 D1-5 模板注入了 UUID 列、或把集中登记同步链拆掉，
`single_html` 的依据就不再成立，但 JSON 文件本身不会自动变红。本判据把四条依据钉成测试：
任一条不再成立 ⇒ 必红并提示回去重新裁决。
"""
from __future__ import annotations

import io
import json
import os
import sys
import warnings
from pathlib import Path

import pytest

_REPO = Path(__file__).resolve().parents[3]
_BACKEND = _REPO / "backend"
if str(_BACKEND) not in sys.path:  # pragma: no cover
    sys.path.insert(0, str(_BACKEND))
os.environ.setdefault("DB_DISABLE_SSL", "True")
warnings.filterwarnings("ignore")

from app.services.workpaper_sync import phase5_d1_notes_receivable as D1  # noqa: E402

MANAGED_SHEET_D105 = "调整分录汇总表D1-5"
_FRONTEND = _REPO / "audit-platform" / "frontend" / "src"
_VERDICT_PATH = (
    _REPO
    / ".kiro/specs/d1-sync-row-table-engine-and-d1-coverage/evidence"
    / "task27-d1-5-single-html-adjudication.json"
)


@pytest.fixture(scope="module")
def verdict() -> dict:
    return json.loads(_VERDICT_PATH.read_text(encoding="utf-8"))


@pytest.fixture(scope="module")
def sheet():
    import openpyxl

    wb = openpyxl.load_workbook(
        io.BytesIO(D1.read_authoritative_template()), data_only=False
    )
    return wb[MANAGED_SHEET_D105]


class TestDecisionIsStillSingleHtml:
    def test_verdict_file_says_single_html(self, verdict: dict) -> None:
        assert verdict["decision"] == "single_html"
        assert len(verdict["why_single_html"]) == 4, "四条依据缺一即裁决不完整"


class TestCriterion1NoRowIdentityColumn:
    """依据①：无行身份列（无 UUID 锚 ⇒ 单元格双向回写无法稳定定位）。"""

    def test_no_uuid_or_gtrow_trace_anywhere(self, sheet) -> None:
        hits = [
            (cell.coordinate, cell.value)
            for row in sheet.iter_rows()
            for cell in row
            if isinstance(cell.value, str)
            and any(t in cell.value.upper() for t in ("GTROW", "_GT", "UUID"))
        ]
        assert not hits, (
            f"D1-5 模板出现行身份列痕迹 {hits} —— 若确实注入了 UUID 列，"
            "single_html 裁决的依据①不再成立，须回去重新裁决"
        )

    def test_d1_3_has_identity_column_as_contrast(self) -> None:
        """反面对照：真正受管的 D1-3 **有** `uuid_col`，证明本判据能区分两者。"""
        spec = D1.instrumentation_spec()
        assert str(spec.managed_sheet) != MANAGED_SHEET_D105
        assert str(spec.uuid_col).strip(), (
            "D1-3 应有非空 uuid_col 作行身份锚——若它也空了，"
            "说明本判据的对照前提失效"
        )


class TestCriterion2DedicatedSyncChainExists:
    """依据②：已有专用集中登记同步链，OO 回写会争同一 store。"""

    def test_d1_tab_adjustment_wires_central_sync_twice(self) -> None:
        """两个**接线点**：1 条 import 语句 + 1 处调用。

        🔴 不断言裸字符串出现次数（实测 = 3，因为 import 那行同时写了符号名与模块路径
        `from '../composables/useAdjustmentCentralSync'`）——直接断言 3 会把「接线点数」
        与「字符出现数」混为一谈，下次有人改 import 写法就误红。
        """
        src = (
            _FRONTEND / "components/workpaper/d1/D1TabAdjustment.vue"
        ).read_text(encoding="utf-8")
        import_lines = [
            ln
            for ln in src.splitlines()
            if ln.strip().startswith("import") and "useAdjustmentCentralSync" in ln
        ]
        assert len(import_lines) == 1, (
            f"预期恰 1 条 import 接线，实得 {len(import_lines)} 条: {import_lines}"
        )
        assert src.count("useAdjustmentCentralSync({") == 1, (
            "预期恰 1 处 useAdjustmentCentralSync({...}) 调用接线 —— "
            "接线点数变了，依据②须复核"
        )
        assert "itemId: 'D1-entry-rows'" in src
        assert "wpCode: 'D1'" in src

    def test_store_key_is_owned_by_use_d1_adjustment(self) -> None:
        src = (
            _FRONTEND / "components/workpaper/composables/useD1Adjustment.ts"
        ).read_text(encoding="utf-8")
        assert "const STORAGE_KEY = 'D1-entry-rows'" in src

    def test_progress_keys_read_the_same_store(self) -> None:
        src = (_FRONTEND / "components/workpaper/d1/D1TabIndex.vue").read_text(
            encoding="utf-8"
        )
        assert "'D1-entry-rows'" in src, (
            "D1TabIndex 的 progressKeys 不再读 D1-entry-rows —— "
            "hub store 的消费面变了，依据②须复核"
        )


class TestCriterion3BalanceEnforcedOnHtmlSideOnly:
    """依据③：借贷平衡仅 HTML 侧强制（Excel 模板零校验）。"""

    def test_html_side_has_balance_tolerance_and_gate(self) -> None:
        composable = (
            _FRONTEND / "components/workpaper/composables/useD1Adjustment.ts"
        ).read_text(encoding="utf-8")
        assert "const BALANCE_TOLERANCE = 0.005" in composable
        assert "isBalanced" in composable
        host = (
            _FRONTEND / "components/workpaper/d1/D1TabAdjustment.vue"
        ).read_text(encoding="utf-8")
        assert "!isBalanced" in host, (
            "「同步到集中登记」按钮不再以 isBalanced 为门 —— 依据③须复核"
        )

    def test_excel_data_band_has_no_formula_and_no_content(self, sheet) -> None:
        """数据区 6-20 是空模板带：零内容零公式 ⇒ Excel 侧不可能校验平衡。"""
        non_empty = [
            (sheet.cell(row=r, column=c).coordinate, sheet.cell(row=r, column=c).value)
            for r in range(6, 21)
            for c in range(1, 11)
            if sheet.cell(row=r, column=c).value is not None
        ]
        assert not non_empty, (
            f"D1-5 数据区 6-20 出现内容/公式 {non_empty[:5]} —— "
            "若模板加了平衡校验公式，依据③不再成立"
        )


class TestSamePatternAsD44Precedent:
    """D1-5 与已判 single_html 的 D4-4 是同一套模板范式（逐字相同的十列表头）。"""

    def test_headers_match_d4_4_verdict_verbatim(self, sheet, verdict: dict) -> None:
        d44_verdict = json.loads(
            (
                _REPO
                / ".kiro/specs/d-cycle-sheet-bidirectional-expansion/evidence"
                / "T08-d44-single-html-adjudication.json"
            ).read_text(encoding="utf-8")
        )
        actual = [sheet.cell(row=5, column=c).value for c in range(1, 11)]
        assert actual == d44_verdict["authority"]["header_A_to_J"], (
            f"D1-5 表头与 D4-4 不再逐字相同：\nD1-5={actual}\n"
            f"D4-4={d44_verdict['authority']['header_A_to_J']}\n"
            "「同一套模板范式」这条类推依据须复核"
        )
        # 本 spec 自己的裁决文件也必须记着同一份表头（两处不得分叉）。
        assert actual == verdict["authority"]["header_A_to_J"]

    def test_both_verdicts_agree_on_single_html(self, verdict: dict) -> None:
        d44_verdict = json.loads(
            (
                _REPO
                / ".kiro/specs/d-cycle-sheet-bidirectional-expansion/evidence"
                / "T08-d44-single-html-adjudication.json"
            ).read_text(encoding="utf-8")
        )
        assert d44_verdict["decision"] == verdict["decision"] == "single_html"

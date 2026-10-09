# -*- coding: utf-8 -*-
"""L2 应付利息真双向接线守卫（spec `l-cycle-true-adapter-registration` · Task 12 最后一条）。

受管表 `应付利息检查表L2-4`，是 L3-9 的孪生凭证检查表，与 L3 同序：
几何现算 → 契约 → 注册 → HTML 侧对齐 → 真库干净命名空间。

L2-4 与 L3-9 的差异（照抄 L3 会错）：
* 数据区 **r12~r23**（12 行，比 L3-9 的 r12~r20 多 3 行），footer **A24**（L3 是 A21）。
* UUID 列 **T**（有效列 P / 物理 max_column S=19），L3 是 AI（max_column AH）。
* 整册 112 裸 IF / 0 definedName（L3 是 38/31 broken + L3-7 四个 #REF!）。
"""
from __future__ import annotations

import re

import pytest
from openpyxl import load_workbook

from app.services.workpaper_sync import phase5_l2_interest_payable as L2
from app.services.workpaper_sync import phase5_l_cycle_common as L
from app.services.workpaper_sync.contracts import parse_contract
from app.services.workpaper_sync.definitions import canonical_digest
from tests.workpaper_sync.l1_adapter_facts import ROOT, _count_bare_if, _is_formula, _query

TEMPLATE = ROOT / "backend" / "wp_templates" / L2.TEMPLATE_RELATIVE_PATH
FRONTEND = ROOT / "audit-platform" / "frontend" / "src" / "components" / "workpaper"
ROWS_TS = FRONTEND / "composables" / "useL2VoucherCheck.ts"


@pytest.fixture(scope="module")
def wb():
    return load_workbook(TEMPLATE, data_only=False)


@pytest.fixture(scope="module")
def ws(wb):
    return wb[L2.MANAGED_SHEET]


# ═══════════════════════════════════════════════════════════════════════════
# 选型与几何
# ═══════════════════════════════════════════════════════════════════════════


class TestManagedSheetSelection:
    def test_template_sha256_matches_frozen(self) -> None:
        import hashlib

        assert hashlib.sha256(TEMPLATE.read_bytes()).hexdigest() == L2.TEMPLATE_SHA256

    def test_managed_sheet_present_and_rejected_candidates_shape(self, wb) -> None:
        """L2-1 是派生审定表（公式远多于文本）—— 选型排除理由可复算。"""
        assert L2.MANAGED_SHEET in wb.sheetnames
        s1 = wb["审定表L2-1"]
        formulas = sum(_is_formula(c.value) for r in s1.iter_rows() for c in r)
        assert formulas >= 20, formulas

    def test_geometry(self, ws) -> None:
        assert ws.max_column == 19, "物理 max_column=S ⇒ UUID 列 T"
        assert L2.UUID_COL == "T"
        assert ws[f"A{L2.FOOTER_ROW}"].value == L2.FOOTER_MARKER
        assert (L2.HEADER_GROUP_ROW, L2.HEADER_LEAF_ROW) == (10, 11)
        assert (L2.FIRST_DATA_ROW, L2.LAST_DATA_ROW, L2.FOOTER_ROW) == (12, 23, 24)

    def test_data_region_has_no_formula_columns(self, ws) -> None:
        """L2-4 受管区零公式（与 L3-9 同，与 L4 的 10 公式列相反）。"""
        from openpyxl.utils import get_column_letter

        for r in range(L2.FIRST_DATA_ROW, L2.LAST_DATA_ROW + 1):
            cols = {
                get_column_letter(c)
                for c in range(1, 17)  # A..P 受管面
                if _is_formula(ws.cell(r, c).value)
            }
            assert cols == set(), (r, sorted(cols))
        assert L2.FORMULA_COLUMNS == ()
        assert L2.FORMULA_MASK == ()

    def test_managed_sheet_has_no_bare_if_but_workbook_does(self, wb, ws) -> None:
        """GC-2 per-file 中性化依据：受管表 0、整册 > 0（实测 112）。"""
        assert _count_bare_if(ws) == 0
        assert sum(_count_bare_if(wb[n]) for n in wb.sheetnames) > 0

    def test_sixteen_managed_fields_are_all_editable(self) -> None:
        assert len(L2.MANAGED_FIELD_SPECS_7) == 16
        assert all(mode == "editable" for _, _, mode, *_ in L2.MANAGED_FIELD_SPECS_7)
        cols = [col for _, col, *_ in L2.MANAGED_FIELD_SPECS_7]
        assert cols == list("ABCDEFGHIJKLMNOP")


# ═══════════════════════════════════════════════════════════════════════════
# 契约
# ═══════════════════════════════════════════════════════════════════════════


class TestContract:
    def test_disk_and_source_are_locked(self) -> None:
        contract = L2.assert_contract_file_matches_source()
        assert contract.canonical_sha256 == canonical_digest(L2.build_contract_payload())

    def test_payload_parses_and_has_expected_shape(self) -> None:
        p = L2.build_contract_payload()
        parse_contract(dict(p), adapter_id=L2.ADAPTER_ID)
        assert p["review_status"] == "reviewed"
        assert p["review"]["entry_id"] == L2.ENTRY_ID
        table = p["sheets"][0]["tables"][0]
        assert table["header_rows"] == 2 and table["anchor"] == "A10"
        assert table["footer_anchor"]["search_column"] == "A"
        assert len(table["fields"]) == 16
        assert table["formula_mask"] == []
        assert not any(f["mode"] == "formula" for f in table["fields"])

    def test_derived_readonly_sheet_declared(self) -> None:
        review = L2.build_contract_payload()["review"]
        drs = review["derived_readonly_sheet"]
        assert drs["excel_name"] == L2.MANAGED_SHEET
        assert "R25" in drs["reason"] and "R26" in drs["reason"]

    def test_routes_through_common_skeleton(self) -> None:
        assert L2.EntrySelectionError is L.LEntrySelectionError
        assert L2.build_contract_payload() == L.build_contract_payload(L2.IDENTITY, L2.SPECS)


# ═══════════════════════════════════════════════════════════════════════════
# 注册
# ═══════════════════════════════════════════════════════════════════════════


class TestRegistration:
    def test_provider_aliases_present(self) -> None:
        assert L2.publish_pilot_definitions is L2.publish_definitions
        assert L2.attach_pilot_adapters is L2.attach_adapters
        assert L2.PILOT_WP_CODES == frozenset({"L2I"})

    def test_store_registry_entry(self) -> None:
        from app.services.workpaper_sync.store_item_registry import STORE_MERGE_REGISTRY

        plan = STORE_MERGE_REGISTRY["l2.interest_payable"]
        assert plan.provider_module == "phase5_l2_interest_payable"
        assert [i.item_id for i in plan.items] == ["L2-L2-4-voucher-rows"]
        assert plan.oo_crash_neutralization_fn == "neutralize_oo_crash_if_formulas"

    def test_whitelist_and_ledger_are_paired(self) -> None:
        from app.services.workpaper_sync.adapters import registry as R
        from app.services.workpaper_sync.adapters.delivered_contracts_ledger import (
            DELIVERED_PER_ENTRY_CONTRACTS,
        )

        mod = "app.services.workpaper_sync.phase5_l2_interest_payable"
        assert mod in R._ALLOWED_PROVIDER_MODULES
        rows = [r for r in DELIVERED_PER_ENTRY_CONTRACTS if r["contract_id"] == "l2.interest_payable"]
        assert len(rows) == 1 and rows[0]["provider_module"] == mod

    def test_wp_code_adjudication_targets_the_whole_workbook_code(self) -> None:
        import json

        doc = json.loads(
            (ROOT / "backend/data/workpaper_sync_entry_wp_code_adjudication.json").read_text("utf-8")
        )
        row = next(a for a in doc["adjudications"] if a["entry_id"] == L2.ENTRY_ID)
        assert row["wp_codes"] == ["L2"] and row["contract_id"] == L2.ADAPTER_ID


# ═══════════════════════════════════════════════════════════════════════════
# HTML 侧对齐
# ═══════════════════════════════════════════════════════════════════════════


class TestHtmlAlignment:
    def test_store_item_id_agrees(self) -> None:
        src = ROWS_TS.read_text("utf-8")
        assert f"L2_VOUCHER_ROWS_ITEM_ID = '{L2.STORE_ITEM_ID}'" in src

    def test_row_identity_uses_shared_minter_not_private_generator(self) -> None:
        """私有 genRowId 已删，改走平台共享 newRowIdentity('l24vc')，已有 rowId 不重铸。"""
        src = ROWS_TS.read_text("utf-8")
        assert "function genRowId" not in src
        assert "newRowIdentity('l24vc')" in src
        assert "raw.rowId || newRowIdentity('l24vc')" in src

    def test_frontend_row_keys_cover_every_contract_field(self) -> None:
        src = ROWS_TS.read_text("utf-8")
        iface = src.split("export interface VoucherCheckRow {", 1)[1].split("}", 1)[0]
        frontend = set(re.findall(r"^\s+(\w+)\s*:", iface, re.M))
        contract = {
            f["json_pointer"].rsplit("/", 1)[-1]
            for f in L2.build_contract_payload()["sheets"][0]["tables"][0]["fields"]
        }
        assert contract - frontend == set(), f"契约有、前端缺：{sorted(contract - frontend)}"


# ═══════════════════════════════════════════════════════════════════════════
# 真库：干净命名空间
# ═══════════════════════════════════════════════════════════════════════════


class TestRealDbNamespace:
    def test_new_voucher_key_and_old_positional_key_are_clean(self) -> None:
        rows = _query(
            "SELECT count(*) FILTER (WHERE item_id = 'L2-L2-4-voucher-rows'), "
            "count(*) FILTER (WHERE item_id LIKE 'L2-L2-4-%-data') FROM checklist_responses"
        )
        new, old_positional = rows[0]
        assert old_positional == 0, "L2-4 旧位置化键有真数据 ⇒ 切换前须先迁移"
        assert new >= 0

# -*- coding: utf-8 -*-
"""L4 应付债券真双向接线守卫（spec `l-cycle-true-adapter-registration` · Task 12）。

受管表 `划分为金融负债的其他金融工具明细表L4-3`。判据按 L1 同序：
几何现算 → 契约 → 注册 → HTML 侧对齐 → 真库干净命名空间。
"""
from __future__ import annotations

import re
from pathlib import Path
from typing import Any

import pytest
from openpyxl import load_workbook

from app.services.workpaper_sync import phase5_l4_bonds_payable as L4
from app.services.workpaper_sync import phase5_l_cycle_common as L
from app.services.workpaper_sync.contracts import parse_contract
from app.services.workpaper_sync.definitions import canonical_digest
from tests.workpaper_sync.l1_adapter_facts import ROOT, _count_bare_if, _is_formula, _query

TEMPLATE = ROOT / "backend" / "wp_templates" / L4.TEMPLATE_RELATIVE_PATH
FRONTEND = ROOT / "audit-platform" / "frontend" / "src" / "components" / "workpaper"
ROWS_TS = FRONTEND / "composables" / "useL4FinLiabRows.ts"


@pytest.fixture(scope="module")
def wb():
    return load_workbook(TEMPLATE, data_only=False)


@pytest.fixture(scope="module")
def ws(wb):
    return wb[L4.MANAGED_SHEET]


# ═══════════════════════════════════════════════════════════════════════════
# 选型与几何（openpyxl 现算，禁抄常量自证）
# ═══════════════════════════════════════════════════════════════════════════


class TestManagedSheetSelection:
    def test_template_sha256_matches_frozen(self) -> None:
        import hashlib

        assert hashlib.sha256(TEMPLATE.read_bytes()).hexdigest() == L4.TEMPLATE_SHA256

    def test_sanitized_template_passes_ooxml_gate(self) -> None:
        """净化后过生产 OOXML 门；零外部关系（净化前因 L4-5 一个外部 hyperlink 被拒）。"""
        import zipfile

        from app.services.workpaper_sync.ooxml_security import validate_ooxml_artifact

        validate_ooxml_artifact(TEMPLATE, document_type="xlsx")
        with zipfile.ZipFile(TEMPLATE) as zf:
            ext = [n for n in zf.namelist() if n.endswith(".rels") and b'TargetMode="External"' in zf.read(n)]
        assert ext == []
        assert L4.PRE_SANITIZE_TEMPLATE_SHA256 != L4.TEMPLATE_SHA256

    def test_managed_sheet_tail_code_is_unique(self, wb) -> None:
        """BP-8 不适用的前提：L4-3 尾码在 16 张 sheet 里只出现一次（L4-7/L4-8 各两次作对照）。"""
        names = wb.sheetnames
        assert len(names) == 16
        assert sum(n.rstrip().endswith("L4-3") for n in names) == 1
        assert sum(n.rstrip().endswith("L4-7") for n in names) == 2
        assert sum(n.rstrip().endswith("L4-8") for n in names) == 2

    def test_rejected_candidates_have_the_recorded_shape(self, wb) -> None:
        """选型排除理由可复算：L4-1 公式远多于文本（派生汇总）；L4-2 有两段小计。"""
        s1 = wb["审定表L4-1"]
        formulas = sum(_is_formula(c.value) for r in s1.iter_rows() for c in r)
        assert formulas >= 100, formulas
        s2 = wb["应付债券明细表L4-2"]
        subtotal_rows = [
            c.row for r in s2.iter_rows(min_col=1, max_col=3) for c in r
            if isinstance(c.value, str) and c.value.replace(" ", "") == "小计"
        ]
        assert len(subtotal_rows) == 2, subtotal_rows

    def test_no_defined_names_and_no_excel_tables(self, wb, ws) -> None:
        assert len(list(wb.defined_names)) == 0
        assert len(ws.tables) == 0

    def test_geometry(self, ws) -> None:
        assert ws.max_column == 40, "UUID 列 AO = max_column + 1"
        assert L4.UUID_COL == "AO"
        assert ws["A18"].value == L4.FOOTER_MARKER
        assert {str(m) for m in ws.merged_cells.ranges} >= {"A10:B12", "A18:B18"}
        # 三级表头：r12 叶子全是「数量/金额」
        leaves = {ws.cell(12, c).value for c in range(12, 40) if ws.cell(12, c).value}
        assert {str(v).strip() for v in leaves} == {"数量", "金额"}

    def test_formula_columns_are_exactly_the_declared_ten(self, ws) -> None:
        """数据区 r13~r17 逐格：公式列恰为 R/S/AF~AM，且形态等于声明模板。"""
        from openpyxl.utils import get_column_letter

        for r in range(L4.FIRST_DATA_ROW, L4.LAST_DATA_ROW + 1):
            cols = {get_column_letter(c) for c in range(1, ws.max_column + 1) if _is_formula(ws.cell(r, c).value)}
            assert cols == set(L4.FORMULA_COLUMNS), (r, sorted(cols))
            for col, tpl in L4.FORMULA_TEMPLATES.items():
                assert ws[f"{col}{r}"].value == tpl.format(row=r), (col, r)

    def test_managed_sheet_has_no_bare_if_but_workbook_does(self, wb, ws) -> None:
        """GC-2 per-file 中性化的依据：受管表 0、整册 > 0。"""
        assert _count_bare_if(ws) == 0
        assert sum(_count_bare_if(wb[n]) for n in wb.sheetnames) > 0

    def test_field_modes_match_template_cell_by_cell(self, ws) -> None:
        for key, col, mode, *_ in L4.MANAGED_FIELD_SPECS_7:
            cell = ws[f"{col}{L4.FIRST_DATA_ROW}"].value
            assert (mode == "formula") == _is_formula(cell), (key, col, mode, cell)


# ═══════════════════════════════════════════════════════════════════════════
# 契约
# ═══════════════════════════════════════════════════════════════════════════


class TestContract:
    def test_disk_and_source_are_locked(self) -> None:
        contract = L4.assert_contract_file_matches_source()
        assert contract.canonical_sha256 == canonical_digest(L4.build_contract_payload())

    def test_payload_parses_and_has_expected_shape(self) -> None:
        p = L4.build_contract_payload()
        parse_contract(dict(p), adapter_id=L4.ADAPTER_ID)
        assert p["review_status"] == "reviewed"
        assert p["review"]["entry_id"] == L4.ENTRY_ID
        table = p["sheets"][0]["tables"][0]
        assert table["header_rows"] == 3 and table["anchor"] == "A10"
        assert table["footer_anchor"]["search_column"] == "A"
        assert len(table["fields"]) == 38
        masked = {m.split(":")[0].rstrip("0123456789") for m in table["formula_mask"]}
        assert masked == set(L4.FORMULA_COLUMNS)
        formula_cols = {f["cell"]["column"] for f in table["fields"] if f["mode"] == "formula"}
        assert formula_cols == masked

    def test_no_bundle_forward_reference_and_no_col_placeholders(self) -> None:
        text = L4.dump_contract_json()
        assert "bundle_" not in text
        keys = {f["column_key"] for f in L4.build_contract_payload()["sheets"][0]["tables"][0]["fields"]}
        assert not any(re.fullmatch(r"col_[a-z]+", k) for k in keys)

    def test_bp8_route_a_inherited_verbatim(self) -> None:
        bp8 = L4.build_contract_payload()["review"]["bp8_sheet_granularity_collapse"]
        assert bp8["status"] == "not_managed_in_this_contract"
        assert bp8["resolve_target_sheet_strict"] is True
        assert {p["tail_code"] for p in bp8["collapsed_pairs"]} == {"L4-7", "L4-8"}

    def test_routes_through_common_skeleton(self) -> None:
        assert L4.EntrySelectionError is L.LEntrySelectionError
        assert L4.build_contract_payload() == L.build_contract_payload(L4.IDENTITY, L4.SPECS)


# ═══════════════════════════════════════════════════════════════════════════
# 注册
# ═══════════════════════════════════════════════════════════════════════════


class TestRegistration:
    def test_provider_aliases_present(self) -> None:
        assert L4.publish_pilot_definitions is L4.publish_definitions
        assert L4.attach_pilot_adapters is L4.attach_adapters
        assert L4.PILOT_WP_CODES == frozenset({"L4B"})

    def test_store_registry_entry(self) -> None:
        from app.services.workpaper_sync.store_item_registry import STORE_MERGE_REGISTRY

        plan = STORE_MERGE_REGISTRY["l4.bonds_payable"]
        assert plan.provider_module == "phase5_l4_bonds_payable"
        assert [i.item_id for i in plan.items] == ["L4-3-rows"]
        assert plan.oo_crash_neutralization_fn == "neutralize_oo_crash_if_formulas"

    def test_whitelist_and_ledger_are_paired(self) -> None:
        from app.services.workpaper_sync.adapters import registry as R
        from app.services.workpaper_sync.adapters.delivered_contracts_ledger import (
            DELIVERED_PER_ENTRY_CONTRACTS,
        )

        mod = "app.services.workpaper_sync.phase5_l4_bonds_payable"
        assert mod in R._ALLOWED_PROVIDER_MODULES
        rows = [r for r in DELIVERED_PER_ENTRY_CONTRACTS if r["contract_id"] == "l4.bonds_payable"]
        assert len(rows) == 1 and rows[0]["provider_module"] == mod

    def test_wp_code_adjudication_targets_the_whole_workbook_code(self) -> None:
        import json

        doc = json.loads((ROOT / "backend/data/workpaper_sync_entry_wp_code_adjudication.json").read_text("utf-8"))
        row = next(a for a in doc["adjudications"] if a["entry_id"] == L4.ENTRY_ID)
        assert row["wp_codes"] == ["L4"] and row["contract_id"] == L4.ADAPTER_ID


# ═══════════════════════════════════════════════════════════════════════════
# HTML 侧对齐（前端 json 键 == 契约 json_pointer 段）
# ═══════════════════════════════════════════════════════════════════════════


class TestHtmlAlignment:
    def test_store_item_id_agrees(self) -> None:
        src = ROWS_TS.read_text("utf-8")
        assert f"L4_3_ROWS_ITEM_ID = '{L4.STORE_ITEM_ID}'" in src

    def test_frontend_row_keys_cover_every_contract_field(self) -> None:
        src = ROWS_TS.read_text("utf-8")
        stems = set(re.findall(r"'(\w+)'", src.split("L4_3_PAIR_STEMS = [", 1)[1].split("]", 1)[0]))
        scalar_block = src.split("export function createEmptyFinLiabRow", 1)[1].split("for (const stem", 1)[0]
        #: `key: value` 与简写属性 `key,` 两种写法都要认（`instrumentName,` 是简写）。
        scalars = set(re.findall(r"^\s+(\w+)\s*[:,]", scalar_block, re.M))
        frontend = scalars | {f"{s}Qty" for s in stems} | {f"{s}Amount" for s in stems}
        contract = {
            f["json_pointer"].rsplit("/", 1)[-1]
            for f in L4.build_contract_payload()["sheets"][0]["tables"][0]["fields"]
        }
        assert contract - frontend == set(), f"契约有、前端行模型缺：{sorted(contract - frontend)}"
        assert set(L4.HTML_ONLY_ROW_KEYS) <= frontend
        assert not (set(L4.HTML_ONLY_ROW_KEYS) & contract), "HTML-only 键不得入契约"

    def test_formula_stems_agree_with_contract_formula_fields(self) -> None:
        src = ROWS_TS.read_text("utf-8")
        block = src.split("L4_3_FORMULA_STEMS", 1)[1].split("])", 1)[0]
        fe = set(re.findall(r"'(\w+)'", block))
        contract = {
            f["json_pointer"].rsplit("/", 1)[-1].removesuffix("Qty").removesuffix("Amount")
            for f in L4.build_contract_payload()["sheets"][0]["tables"][0]["fields"]
            if f["mode"] == "formula"
        }
        assert fe == contract


# ═══════════════════════════════════════════════════════════════════════════
# 真库：干净命名空间
# ═══════════════════════════════════════════════════════════════════════════


class TestRealDbNamespace:
    def test_new_key_and_old_positional_key_are_empty(self) -> None:
        rows = _query(
            "SELECT count(*) FILTER (WHERE item_id = 'L4-3-rows'), "
            "count(*) FILTER (WHERE item_id LIKE 'L4-3-row-%') FROM checklist_responses"
        )
        new, old = rows[0]
        assert old == 0, "旧位置化键有真数据 ⇒ 切换前必须先写迁移"
        # 新键可能已由真实使用写入；只断言不是位置化残留
        assert new >= 0

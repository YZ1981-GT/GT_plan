# -*- coding: utf-8 -*-
"""D 循环检查表 / 关联方表：契约逐列表头 == 权威模板逐格（spec oo-writeback-template-parity-closure R3 批 B2a）。

FC-5「以模板为权威」此前只在各 lane 的同构测试里逐家钉，这几张没有：D7-7 两区整表错位（A..R 而模板 A..Q、
多出第 6/7 个检查列、索引号错一列）、D6-6 区② 照抄区① 的列（贷方金额写进支持性文件列）、D6-5 J~N 五列与模板
不符，都是逐格对表头就能看出来的错。判据（逐字段）：
  ① 叶子表头格（两级表头取叶子行；合并格取左上角）归一后 == `header_text`；叶子是编号（核对内容下的 1~5）时，
     `header_text` == 组标题 + 编号；
  ② 声明了组标题格的字段：组标题格的文字是该组标签，且字段列落在组标题的合并区里（列错位必红）；
  ③ 首个数据行紧接表头，footer 行 A 列是 footer marker。
"""
from __future__ import annotations

import os
import re
import sys
from pathlib import Path

import openpyxl
import pytest
from openpyxl.utils import column_index_from_string

_BACKEND = Path(__file__).resolve().parents[2]
if str(_BACKEND) not in sys.path:  # pragma: no cover
    sys.path.insert(0, str(_BACKEND))
os.environ.setdefault("DB_DISABLE_SSL", "True")

from app.services.workpaper_sync.phase5_d1_13_sampling import SPEC_D113_SPECIFIC, SPEC_D113_VOUCHING  # noqa: E402
from app.services.workpaper_sync.phase5_d3_07_voucher_check import SPEC_D307_CURRENT, SPEC_D307_POST  # noqa: E402
from app.services.workpaper_sync.phase5_d6_05_related_party import SPEC_D605  # noqa: E402
from app.services.workpaper_sync.phase5_d6_06_inspection import SPEC_D606_BLOCK1, SPEC_D606_BLOCK2  # noqa: E402
from app.services.workpaper_sync.phase5_d7_07_inspection import SPEC_D707_PERIOD, SPEC_D707_POST  # noqa: E402
from app.services.workpaper_sync.phase5_row_table_sheet import RowTableSheetSpec, managed_field_specs  # noqa: E402

_TPL = _BACKEND / "wp_templates" / "D"
CASES: list[tuple[str, str, RowTableSheetSpec]] = [
    ("D1-13 区①", "D1 应收票据.xlsx", SPEC_D113_VOUCHING),
    ("D1-13 区②", "D1 应收票据.xlsx", SPEC_D113_SPECIFIC),
    ("D3-7 区①", "D3 预收账款.xlsx", SPEC_D307_CURRENT),
    ("D3-7 区②", "D3 预收账款.xlsx", SPEC_D307_POST),
    ("D6-5", "D6 合同资产.xlsx", SPEC_D605),
    ("D6-6 区①", "D6 合同资产.xlsx", SPEC_D606_BLOCK1),
    ("D6-6 区②", "D6 合同资产.xlsx", SPEC_D606_BLOCK2),
    ("D7-7 区①", "D7 合同负债.xlsx", SPEC_D707_PERIOD),
    ("D7-7 区②", "D7 合同负债.xlsx", SPEC_D707_POST),
]
_BOOKS: dict[str, object] = {}


def _norm(v: object) -> str:
    t = re.sub(r"[\s\u3000]+", "", "" if v is None else str(v))
    return t.replace("（", "(").replace("）", ")")


def _sheet(book: str, name: str):
    if book not in _BOOKS:
        _BOOKS[book] = openpyxl.load_workbook(_TPL / book, data_only=False)
    return _BOOKS[book][name]


def _merged_range(ws, row: int, col: int):
    for mr in ws.merged_cells.ranges:
        if mr.min_row <= row <= mr.max_row and mr.min_col <= col <= mr.max_col:
            return mr
    return None


def _value(ws, row: int, col: int) -> object:
    mr = _merged_range(ws, row, col)
    return ws.cell(row=mr.min_row, column=mr.min_col).value if mr else ws.cell(row=row, column=col).value


def header_problems(ws, spec: RowTableSheetSpec) -> list[str]:
    """判据本体（变异自检也走它）。"""
    leaf = spec.header_leaf_row or spec.header_row
    out: list[str] = []
    if leaf is None:
        return ["没有声明表头行"]
    if spec.first_data_row != leaf + 1:
        out.append(f"首个数据行 {spec.first_data_row} 不紧接表头 {leaf}")
    if _norm(ws[f"{spec.footer_search_column}{spec.footer_row}"].value) != _norm(spec.footer_marker):
        out.append(f"footer {spec.footer_row} 不是「{spec.footer_marker}」")
    for _key, col, _mode, _vt, json_path, header_text, group in managed_field_specs(spec):
        c = column_index_from_string(col)
        cell = _norm(_value(ws, leaf, c))
        group_text = _norm(ws[group].value) if group else ""
        if cell.isdigit() and group_text:
            ok = _norm(header_text) == group_text + cell
        else:
            ok = cell == _norm(header_text)
        if not ok:
            out.append(f"{col}{leaf} 模板「{cell}」≠ 契约「{header_text}」（{json_path}）")
        if group:
            g = ws[group]
            mr = _merged_range(ws, g.row, g.column)
            if not group_text or mr is None or not (mr.min_col <= c <= mr.max_col):
                out.append(f"{col} 不在组标题 {group}「{group_text}」的合并区里")
    return out


@pytest.mark.parametrize("label,book,spec", CASES, ids=[c[0] for c in CASES])
def test_headers_match_template_cell_by_cell(label: str, book: str, spec: RowTableSheetSpec) -> None:
    ws = _sheet(book, spec.managed_sheet)
    assert header_problems(ws, spec) == [], label


def test_managed_columns_stop_at_template_header_width() -> None:
    """D6-5 原声明 M / N 落在模板空列上 —— 受管列不得越过表头最右一个非空格。"""
    for label, book, spec in CASES:
        ws = _sheet(book, spec.managed_sheet)
        leaf = spec.header_leaf_row or spec.header_row
        group = spec.header_group_row or leaf
        last = max(
            c for c in range(1, ws.max_column + 1)
            if any(_value(ws, r, c) not in (None, "") for r in range(group, leaf + 1))
        )
        cols = [column_index_from_string(f[1]) for f in managed_field_specs(spec)]
        assert max(cols) <= last, (label, max(cols), last)


# ═══ 变异自检：判据本体喂改过的声明，必须报 ═══


def test_mutation_shifted_columns_are_reported() -> None:
    """D7-7 区① 改回原错位声明的一段（F 借方 / G 贷方 / H 支持性文件）⇒ 表头与组标题都报。"""
    import dataclasses

    ws = _sheet("D7 合同负债.xlsx", SPEC_D707_PERIOD.managed_sheet)
    shifted = tuple(
        (k, {"G": "F", "H": "G", "I": "H"}.get(c, c) if k in ("debit_amount", "credit_amount", "support_doc") else c,
         m, vt, jp, h, ("" if k == "support_doc" else g))
        for k, c, m, vt, jp, h, g in SPEC_D707_PERIOD.field_specs
        if k != "counter_detail"
    )
    bad = header_problems(ws, dataclasses.replace(SPEC_D707_PERIOD, field_specs=shifted))
    assert any("F17" in b and "借方金额" in b for b in bad), bad
    assert any("H17" in b and "支持性文件" in b for b in bad), bad


def test_mutation_single_header_row_and_extra_column_are_reported() -> None:
    import dataclasses

    ws = _sheet("D6 合同资产.xlsx", SPEC_D605.managed_sheet)
    extra = (*SPEC_D605.field_specs, ("plan", "M", "editable", "text", "plan", "处理计划", ""))
    assert any("M9" in b for b in header_problems(ws, dataclasses.replace(SPEC_D605, field_specs=extra)))
    ws7 = _sheet("D7 合同负债.xlsx", SPEC_D707_POST.managed_sheet)
    single = dataclasses.replace(SPEC_D707_POST, header_group_row=None, header_leaf_row=None, header_row=33)
    assert any("首个数据行" in b for b in header_problems(ws7, single))

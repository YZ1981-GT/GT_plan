# -*- coding: utf-8 -*-
"""G 九条四条红判据：P15 G3 definedName · P16 `-修订前` 逐字排除 · P17 prefill 不回归 · P18 零回归现算。

spec: `g-cycle-single-region-detail-lanes` · Task 6 · Requirements 4.4 / 4.5 / 4.6 / 4.7 / 4.8

🔴 **GC-10 现算铁律**：本文件**不写死任何总数**（digest 条数 / 契约条数 / definedName 条数）。
并发会话正在往契约目录加条目（本轮实测：契约目录 17 文件 = 16 生产契约 + 1 `_example.candidate.json`，
而 golden digest 的 `providers` 只有 9 条 —— **两者本就不同步**）。任何写死的数字明天就假红。
现算值只作为**同一次运行内**的前后比对基准，并写进断言消息供人核。
"""
from __future__ import annotations

import json
import os
import sys
from pathlib import Path

import openpyxl
import pytest

_REPO = Path(__file__).resolve().parents[3]
_BACKEND = _REPO / "backend"
if str(_BACKEND) not in sys.path:
    sys.path.insert(0, str(_BACKEND))
if str(Path(__file__).parent) not in sys.path:
    sys.path.insert(0, str(Path(__file__).parent))
os.environ.setdefault("DB_DISABLE_SSL", "True")

from test_g_single_region_p1_p3_p5_p6 import EXPECTED, LANE_ORDER, TPL_G  # noqa: E402

PREFILL_PATH = _BACKEND / "data" / "prefill_formula_mapping.json"
CONTRACTS_DIR = _BACKEND / "data" / "workpaper_sync_contracts"
GOLDEN_DIGEST = _BACKEND / "scripts" / "check" / "_sync_provider_golden_digest.json"

#: 🔴 四张 `-修订前` hidden sheet 的**逐字**名（实测 2026-09-27，含空格）。
#:   `信用减值损失审计程序表G14A -修订前` 的 `G14A` 后有**一个空格** —— 不得 strip、不得去空格。
REVISED_BEFORE_SHEETS: tuple[tuple[str, str, int, int], ...] = (
    ("G11", "投资收益实质性程序表G11A-修订前", 34, 56),
    ("G12", "净敞口套期收益审计程序表G12A-修订前", 63, 115),
    ("G13", "公允价值变动收益审计程序表G13A-修订前", 46, 80),
    ("G14", "信用减值损失审计程序表G14A -修订前", 46, 78),
)


def defined_names_of(book: str) -> frozenset[str]:
    """现算某册的 workbook 级 definedName 名称集合（P15 的可复算 helper）。

    Task 14 / Task 15 受管后复算并与受管前比对 —— 集合**逐项**相等，不比数量。
    """
    wb = openpyxl.load_workbook(TPL_G / book, data_only=False)
    try:
        return frozenset(wb.defined_names.keys())
    finally:
        wb.close()


# ════════════════════════════════════════════════════════════════════════════
# G1R-P15：G3 的 definedName 集合逐项不变（现算，不写死 491）
#   Validates: Requirements 4.4
# ════════════════════════════════════════════════════════════════════════════
class TestG1rP15G3DefinedNames:
    def test_g3_has_a_large_legacy_defined_name_residue_and_g1_has_none(self) -> None:
        """A 类：G3 带大量 legacy 命名区域残留，其余八条为 0（红基线 B7）。

        🔴 断言**数量级与对比关系**，不断言确切值（spec 说「约 480」，本次现算 491）。
        """
        g3 = defined_names_of(EXPECTED["G3"]["workbook"])
        assert len(g3) > 300, f"G3 的 definedName 现算 {len(g3)} 个 —— 少于 300 说明已被清理过"
        others = {
            code: len(defined_names_of(spec["workbook"]))
            for code, spec in EXPECTED.items()
            if code != "G3"
        }
        assert set(others.values()) == {0}, (
            f"除 G3 外其余八条的 definedName 应全为 0，实得 {others}"
        )

    def test_g3_defined_names_contain_the_three_legacy_families(self) -> None:
        """A 类：三族残留各有命中（`_xlnm.*` / `UFPrn*` / 中文名）——
        证明这批名字**不是**受管可以顺手删掉的技术性名字，而是 legacy 工作簿的业务残留。
        """
        names = defined_names_of(EXPECTED["G3"]["workbook"])
        xlnm = {n for n in names if n.startswith("_xlnm")}
        ufprn = {n for n in names if n.upper().startswith("UFPRN")}
        chinese = {n for n in names if any("\u4e00" <= ch <= "\u9fff" for ch in n)}
        assert xlnm, "无 `_xlnm.*`（Print_Area / Database 等）"
        assert ufprn, "无 `UFPrn*`（用友打印模板残留）"
        assert chinese, "无中文名（业务命名残留）"
        assert len(xlnm | ufprn | chinese) <= len(names)

    def test_defined_names_helper_is_deterministic_and_reusable(self) -> None:
        """P15 的落地方式：helper 两次现算结果相等 ⇒ Task 14/15 可拿它做受管前后比对。

        🔴 判据不能只说「集合不变」而不给可执行的取值方式 —— 那样 Task 14 只能凭感觉。
        """
        a = defined_names_of(EXPECTED["G3"]["workbook"])
        b = defined_names_of(EXPECTED["G3"]["workbook"])
        assert a == b and isinstance(a, frozenset)

    def test_mutation_comparing_by_count_would_miss_a_rename(self) -> None:
        """🔴 变异自检：若 P15 按「数量相等」比对，**改名不改数**会被放过。

        构造一个同量异名集合，断言「数量比对通过而集合比对失败」。
        """
        names = defined_names_of(EXPECTED["G3"]["workbook"])
        one = next(iter(sorted(names)))
        mutated = (names - {one}) | {one + "_renamed"}
        assert len(mutated) == len(names), "变异体必须保持数量不变（否则证不出数量比对的漏洞）"
        assert mutated != names, "集合比对必须能发现改名"


# ════════════════════════════════════════════════════════════════════════════
# G1R-P16：四张 `-修订前` hidden sheet 逐字排除（空格不得 strip）
#   Validates: Requirements 4.5
# ════════════════════════════════════════════════════════════════════════════
class TestG1rP16RevisedBeforeSheetsExcludedVerbatim:
    @pytest.mark.parametrize("code,title,rows,merged", REVISED_BEFORE_SHEETS)
    def test_each_sheet_exists_hidden_with_expected_geometry(
        self, code: str, title: str, rows: int, merged: int
    ) -> None:
        """A 类：四张逐字名存在、`sheet_state == 'hidden'`、行数/合并数与实测一致。"""
        wb = openpyxl.load_workbook(TPL_G / EXPECTED[code]["workbook"], data_only=False)
        try:
            assert title in wb.sheetnames, (
                f"{code} 册里找不到逐字名 {title!r}；实际 sheetnames 含 "
                f"{[s for s in wb.sheetnames if '修订前' in s]}"
            )
            ws = wb[title]
            assert ws.sheet_state == "hidden", f"{title!r} 实得 state={ws.sheet_state}"
            assert ws.max_row == rows and len(ws.merged_cells.ranges) == merged
        finally:
            wb.close()

    def test_g14_title_has_an_inner_space_before_the_dash(self) -> None:
        """🔴 P16 的核心：`信用减值损失审计程序表G14A -修订前` 的 `G14A` 后有**一个空格**。

        它在名字**中间**（不是首尾）⇒ `str.strip()` 去不掉它，但任何「去空格再比」的
        归一化都会把它变成另一个名字。
        """
        title = "信用减值损失审计程序表G14A -修订前"
        assert title == next(t for c, t, _r, _m in REVISED_BEFORE_SHEETS if c == "G14")
        assert " -修订前" in title, "空格必须紧跟在 G14A 之后"
        assert title.strip() == title, "空格在名字中间，strip() 不会改变它"
        assert title.replace(" ", "") != title, "去空格会得到另一个名字"

    def test_mutation_space_stripped_name_is_not_in_the_workbook(self) -> None:
        """🔴 变异自检（spec 指定的那个变异）：排除清单若写「去空格版」，则与真名不匹配 ⇒ 漏排除。

        逐字证明：去空格后的名字**不在** G14 册的 sheetnames 里。
        """
        wb = openpyxl.load_workbook(TPL_G / EXPECTED["G14"]["workbook"], data_only=False)
        try:
            real = "信用减值损失审计程序表G14A -修订前"
            mutated = real.replace(" ", "")
            assert real in wb.sheetnames
            assert mutated not in wb.sheetnames, (
                f"去空格名 {mutated!r} 竟然也在 sheetnames 里 ⇒ 变异无害，本判据要重写"
            )
        finally:
            wb.close()

    def test_the_four_are_all_in_this_specs_four_pl_workbooks(self) -> None:
        """A 类：四张全落在本 spec 的四册损益类里（红基线 B10）—— 不在其余五册。"""
        owners = {c for c, _t, _r, _m in REVISED_BEFORE_SHEETS}
        assert owners == {"G11", "G12", "G13", "G14"}
        for code in set(EXPECTED) - owners:
            wb = openpyxl.load_workbook(TPL_G / EXPECTED[code]["workbook"], data_only=False)
            try:
                assert not [s for s in wb.sheetnames if "修订前" in s], (
                    f"{code} 册也出现 `-修订前` 残留 ⇒ 排除清单要扩"
                )
            finally:
                wb.close()

    def test_g12_revised_before_is_the_largest_of_the_four(self) -> None:
        """A 类：G12 的那张 63 行 × 115 合并是四张里最大（design §顺带发现 3）。"""
        by_merged = {c: m for c, _t, _r, m in REVISED_BEFORE_SHEETS}
        assert max(by_merged, key=by_merged.get) == "G12"
        assert by_merged["G12"] == 115


# ════════════════════════════════════════════════════════════════════════════
# G1R-P17：prefill 两块 sheet 名不回归（foundation Task 7 已修）
#   Validates: Requirements 4.6
# ════════════════════════════════════════════════════════════════════════════
@pytest.fixture(scope="module")
def prefill_blocks() -> list[dict]:
    d = json.loads(PREFILL_PATH.read_text(encoding="utf-8"))
    blocks = d if isinstance(d, list) else (
        d.get("blocks") or d.get("mappings") or next(v for v in d.values() if isinstance(v, list))
    )
    assert isinstance(blocks, list) and blocks, "prefill 结构变了，判据须改写"
    return blocks


class TestG1rP17PrefillSheetNamesDoNotRegress:
    def test_g13_and_g14_main_table_blocks_use_the_real_sheet_names(
        self, prefill_blocks
    ) -> None:
        """A 类：G13/G14 的**主受管表**块 `sheet` 是 foundation 修后的真名。

        🔴 实测纠正：**同一个 `wp_code` 有多个 prefill 块**（审定表 / 明细表 / 附注披露 ×2）。
        首版判据对该 wp_code 的**所有**块断言主表名 ⇒ 被 `审定表G13-1` 打红。
        ⇒ 定位方式改为「该 entry 的 `managed_sheet`」，并断言它**恰好命中一块**。
        """
        for code in ("G13", "G14"):
            want = EXPECTED[code]["managed_sheet"]
            hits = [
                b for b in prefill_blocks
                if b.get("wp_code") == code and b.get("sheet") == want
            ]
            assert len(hits) == 1, (
                f"{code} 的主受管表块（sheet={want!r}）实得 {len(hits)} 个；"
                f"该 wp_code 下现有 sheet="
                f"{[b['sheet'] for b in prefill_blocks if b.get('wp_code') == code]} —— "
                "若主表块消失或改名，就是 foundation Task 7 的修复被回退"
            )
        # 逐字反证：错名版本不得出现在任何块上
        wrong = {"明细分析表G13-2", "明细分析表G14-2"}
        present = {b.get("sheet") for b in prefill_blocks}
        assert not (wrong & present), f"prefill 里出现错名 {sorted(wrong & present)}"

    def test_g11_main_table_block_legitimately_uses_the_analysis_name(
        self, prefill_blocks
    ) -> None:
        """🔴 对照组：G11 的真名**就是**带「分析」的 `明细分析表G11-2` ——
        它不是错名，别在「统一去掉分析二字」时把它一起改掉。
        """
        assert EXPECTED["G11"]["managed_sheet"] == "明细分析表G11-2"
        hits = [
            b for b in prefill_blocks
            if b.get("wp_code") == "G11" and b.get("sheet") == "明细分析表G11-2"
        ]
        assert len(hits) == 1, (
            "G11 的主受管表块应恰有一个且 sheet 带「分析」；该 wp_code 下现有 sheet="
            f"{[b['sheet'] for b in prefill_blocks if b.get('wp_code') == 'G11']}"
        )
        assert "明细表G11-2" not in {b.get("sheet") for b in prefill_blocks}, (
            "出现不带「分析」的 G11 主表名 ⇒ 有人把正确名字「修」错了"
        )

    def test_only_five_of_the_nine_have_a_main_table_prefill_block(
        self, prefill_blocks
    ) -> None:
        """🔴 顺带发现（登记，本 spec 不补）：九条里只有 **G1/G8/G11/G13/G14 五条**
        有主受管表 prefill 块；**G3/G9/G10/G12 四条没有**。

        ⇒ P17 的「两块不回归」只覆盖 G13/G14；另外四条**根本没有**可回归的块。
        这是 prefill 覆盖缺口（不是错名），归模板/取数配置治理，本 spec 只登记。
        """
        have = {
            code for code in LANE_ORDER
            if any(
                b.get("wp_code") == code and b.get("sheet") == EXPECTED[code]["managed_sheet"]
                for b in prefill_blocks
            )
        }
        assert have == {"G1", "G8", "G11", "G13", "G14"}, (
            f"有主表 prefill 块的实得 {sorted(have)} —— 若四条缺口被补上，更新本登记判据"
        )
        assert set(LANE_ORDER) - have == {"G3", "G9", "G10", "G12"}

    def test_every_g_block_sheet_name_exists_in_its_template(self, prefill_blocks) -> None:
        """🔴 比「两块不回归」更强的判据：本 spec 九条的**每个** prefill 块，
        其 `sheet` 名必须真存在于对应模板（GC-8 的 sheet 存在性守卫在九条上的落地）。

        这条能抓住将来任何新增块的错名，不只守着 [169]/[170] 两处。
        """
        by_code = {c: s for c, s in ((c, s) for c, s in EXPECTED.items())}
        checked = 0
        for b in prefill_blocks:
            code = b.get("wp_code")
            if code not in by_code:
                continue
            wb = openpyxl.load_workbook(TPL_G / by_code[code]["workbook"], data_only=False)
            try:
                assert b["sheet"] in wb.sheetnames, (
                    f"prefill 块 wp_code={code} 的 sheet {b['sheet']!r} 不在模板 "
                    f"{by_code[code]['workbook']!r} 里（GC-8 sheet 存在性）"
                )
            finally:
                wb.close()
            checked += 1
        assert checked >= 3, f"本 spec 九条相关的 prefill 块只核到 {checked} 个 —— 少于预期"

    def test_index_based_lookup_agrees_with_code_based_for_now(self, prefill_blocks) -> None:
        """交叉核对：当前下标 `[169]`/`[170]` 确实落在 G13/G14 块上。

        🔴 本条允许将来因并发插块而失效 —— 失效时**改本条**，不要改上面按 wp_code 的判据。
        断言消息里写明这一点，避免后人改错地方。
        """
        if len(prefill_blocks) <= 170:
            pytest.skip(f"prefill 只有 {len(prefill_blocks)} 块，下标交叉核对不适用")
        assert prefill_blocks[169].get("wp_code") == "G13", (
            f"下标 [169] 实得 wp_code={prefill_blocks[169].get('wp_code')!r} —— "
            "块已被重排；按 wp_code 的判据仍有效，只需更新本条的下标"
        )
        assert prefill_blocks[170].get("wp_code") == "G14"

    def test_trailing_space_in_a_sibling_block_is_registered_not_fixed(
        self, prefill_blocks
    ) -> None:
        """🔴 顺带发现（登记，不修）：`[171]` 的 `sheet` 是 `'明细表J1-2 '` —— **尾部带空格**。

        它属 J 循环（本 spec 范围外）。登记在此是为了让 J 循环的 spec 不必重新发现；
        本判据只断言「该形态存在」，不改它。若将来 J 循环修了，本条会红 ⇒ 届时删掉本条。
        """
        j1 = [b for b in prefill_blocks if b.get("wp_code") == "J1"]
        if not j1:
            pytest.skip("prefill 里没有 J1 块")
        trailing = [b["sheet"] for b in j1 if b["sheet"] != b["sheet"].rstrip()]
        assert trailing == ["明细表J1-2 "], (
            f"J1 块的尾空格形态实得 {trailing} —— 若已修好，删掉本条登记判据"
        )


# ════════════════════════════════════════════════════════════════════════════
# G1R-P18：零回归**现算逐项** + G9 不 seed 而其余八条 seed
#   Validates: Requirements 4.7 / 4.8
# ════════════════════════════════════════════════════════════════════════════
class TestG1rP18ZeroRegressionComputedNotHardcoded:
    def test_contract_dir_and_golden_digest_are_not_in_sync_so_counts_must_be_computed(
        self,
    ) -> None:
        """🔴 GC-10 的可执行论据：契约目录条数 ≠ golden digest 的 provider 条数。

        本次现算：契约目录 16 个生产契约（+1 个 `_example.candidate.json`）、digest 9 个 provider。
        ⇒ 任何「断言两者相等」或「写死某个数」的判据都会在并发会话加契约时假红。
        """
        prod_contracts = {
            p.stem for p in CONTRACTS_DIR.glob("*.json") if not p.name.startswith("_")
        }
        digest = json.loads(GOLDEN_DIGEST.read_text(encoding="utf-8"))
        providers = {p["adapter_id"] for p in digest["providers"]}
        assert prod_contracts, "契约目录为空"
        assert providers, "golden digest 无 provider"
        assert prod_contracts != providers, (
            f"契约目录（{len(prod_contracts)}）与 digest（{len(providers)}）竟然一致 —— "
            "若真同步了，本判据的论据消失，但 GC-10 的现算要求仍然有效（不要改成写死）"
        )

    def test_all_nine_adapters_are_in_the_golden_digest(self) -> None:
        """🔴 **B 类红判据（正向）**：九条的 adapter 全部进入 golden digest。

        写成正向（「全在」）而不是倒置（「全不在」）：倒置断言现在是绿的，Task 15 交付后
        才变红并需要人去反转 —— 那等于把「该红的时候不红」制度化了。
        正向红判据在 Task 15 之前一直红，交付后自动转绿，无需改判据。
        """
        digest = json.loads(GOLDEN_DIGEST.read_text(encoding="utf-8"))
        providers = {p["adapter_id"] for p in digest["providers"]}
        nine = {EXPECTED[c]["adapter_id"] for c in LANE_ORDER}
        missing = sorted(nine - providers)
        assert not missing, (
            f"以下九条 adapter 尚未进 golden digest（Task 15 转绿）：{missing}；"
            f"digest 现有 {len(providers)} 条 provider（现算，勿写死）"
        )

    def test_all_nine_contracts_are_published(self) -> None:
        """🔴 **B 类红判据（正向）**：九条的 per-entry 契约文件全部发布。"""
        existing = {p.stem for p in CONTRACTS_DIR.glob("*.json")}
        nine = {EXPECTED[c]["adapter_id"] for c in LANE_ORDER}
        missing = sorted(nine - existing)
        assert not missing, (
            f"以下九条契约尚未发布（Task 15 转绿）：{missing}；"
            f"契约目录现有 {len(existing)} 个文件（现算，勿写死）"
        )

    def test_zero_regression_baseline_is_captured_per_item_not_as_a_count(self) -> None:
        """P18 的落地方式：把既有 digest 逐项快照成 `{adapter_id: (contract_sha, sheet_digests)}`。

        Task 15 复算时**逐项**比对这个映射（本 spec 未触及的条目必须逐字节相等），
        🔴 而不是比 `digest_count`（那个数会被并发会话改）。
        """
        digest = json.loads(GOLDEN_DIGEST.read_text(encoding="utf-8"))
        snapshot = {
            p["adapter_id"]: (
                p.get("contract_payload_sha256"),
                tuple(sorted((p.get("sheet_digests") or {}).items())),
                p.get("store_projection_sha256"),
                p.get("instrumentation_sha256"),
            )
            for p in digest["providers"]
        }
        assert len(snapshot) == len(digest["providers"]), "adapter_id 有重复 ⇒ 快照会丢条目"
        # 每条都要有 contract sha（否则快照不足以发现改动）
        missing = [a for a, v in snapshot.items() if not v[0]]
        assert not missing, f"以下 provider 缺 contract_payload_sha256：{missing}"

    def test_g9_is_the_only_lane_that_can_skip_seeding(self) -> None:
        """Req 4.8：只有 G9 有真库载荷（605 B / 2 wp）可不 seed；其余八条必须 seed。"""
        adj = json.loads(
            (_BACKEND / "data" / "workpaper_sync_entry_wp_code_adjudication.json").read_text(
                encoding="utf-8"
            )
        )
        by_entry = {a["entry_id"]: a for a in adj["adjudications"] if a.get("entry_id")}
        with_payload = {
            code
            for code in LANE_ORDER
            if by_entry[EXPECTED[code]["entry_id"]]["store_payload_evidence"][
                "max_payload_bytes"
            ]
            > 2
        }
        assert with_payload == {"G9"}, (
            f"载荷 >2 B 的条实得 {sorted(with_payload)} —— 若变了，seed 清单要同步改"
        )
        ev = by_entry[EXPECTED["G9"]["entry_id"]]["store_payload_evidence"]
        assert ev["max_payload_bytes"] == 605
        # 🔴 术语纠正：spec 写「605 B / **2 wp**」，实测 `wp_count_with_payload == 1`。
        #    note 逐字：「真库两行：一行 605 B（真实载荷）、一行 2 B（空数组）」
        #    ⇒ 该键在 **2 个 wp 上有 store 行**，但**只有 1 个 wp 的载荷非空**。
        #    两个数不是一回事；seed 判据若按「2 wp 有载荷」写，会以为已有两份真数据。
        assert ev["wp_count_with_payload"] == 1, (
            f"G9 有非空载荷的 wp 数实得 {ev['wp_count_with_payload']} —— "
            "spec 的「2 wp」说的是 store 行数（含 1 行空数组），不是有载荷的 wp 数"
        )
        assert ev["wp_code_with_payload"] == "G9"
        assert "一行 605 B" in ev["note"] and "一行 2 B" in ev["note"]

    def test_g1_and_g3_need_key_creation_not_just_row_insertion(self) -> None:
        """🔴 Task 1 §3.1 的结论落成判据：G1/G3 是「键不存在」（0 B / 0 wp），
        与「有记录但值是 `[]`（2 B）」是两种状态 ⇒ seed 须先**建键**。

        P18 的「未 seed 时验收脚本显式失败」必须能区分这两态，否则 G1/G3 会以「2 B 空数组」
        的假象通过。
        """
        adj = json.loads(
            (_BACKEND / "data" / "workpaper_sync_entry_wp_code_adjudication.json").read_text(
                encoding="utf-8"
            )
        )
        by_entry = {a["entry_id"]: a for a in adj["adjudications"] if a.get("entry_id")}
        states = {
            code: (
                by_entry[EXPECTED[code]["entry_id"]]["store_payload_evidence"][
                    "max_payload_bytes"
                ],
                by_entry[EXPECTED[code]["entry_id"]]["store_payload_evidence"][
                    "wp_count_with_payload"
                ],
            )
            for code in LANE_ORDER
        }
        assert states["G1"] == (0, 0) and states["G3"] == (0, 0), f"实得 {states}"
        two_byte = {c for c, (b, _w) in states.items() if b == 2}
        assert two_byte == {"G8", "G10", "G11", "G12", "G13", "G14"}, (
            f"2 B 空数组的条实得 {sorted(two_byte)}"
        )


# ════════════════════════════════════════════════════════════════════════════
# G1R-P13：九条 sync 路径对 trial_balance 写次数全部为 0
#   Validates: Requirements 4.1, 4.2
# ════════════════════════════════════════════════════════════════════════════
class TestG1rP13TbRedLine:
    """FC-9 红线：九条 sync provider 对 trial_balance 的写次数全部为 0。

    TB 写入由审定表的 `publishToTb` 显式确认门承载（归后置 spec
    `g-cycle-adjudication-sheets-coverage`），不在 sync 路径内。
    """

    def test_no_provider_source_contains_tb_write_logic(self) -> None:
        """九条 provider 的源代码不含 trial_balance 写逻辑。

        允许出现 "trial_balance 写次数为 0" 之类的文档注释（那正是声明），
        但不允许出现 `INSERT INTO trial_balance` / `update.*trial_balance` / `publishToTb` 等写入形态。
        """
        import importlib
        import inspect
        import re

        from app.services.workpaper_sync.store_item_registry import STORE_MERGE_REGISTRY

        # 🔴 GC-10 现算：本判据只覆盖 single-region 九条（从 EXPECTED 的 LANE_ORDER 取），
        #    不写死数量。g4/g5/g6 有 TB 发布门（GC-9 三家 + G5），另由 g4-g6/g5 spec 各自守。
        nine = {EXPECTED[c]["adapter_id"] for c in LANE_ORDER}
        g_adapters = {aid for aid in STORE_MERGE_REGISTRY if aid in nine}
        assert g_adapters == nine, (
            f"single-region 九条 adapter 未全在 registry：缺 {nine - g_adapters}"
        )

        for aid in sorted(g_adapters):
            plan = STORE_MERGE_REGISTRY[aid]
            mod_name = f"app.services.workpaper_sync.{plan.provider_module}"
            mod = importlib.import_module(mod_name)
            source = inspect.getsource(mod)
            # 排除文档字符串（三引号区块）和单行注释
            import ast as _ast
            try:
                tree = _ast.parse(source)
            except SyntaxError:
                continue
            # 收集所有 docstring 的行号范围
            doc_lines: set[int] = set()
            for node in _ast.walk(tree):
                if isinstance(node, (_ast.FunctionDef, _ast.AsyncFunctionDef, _ast.ClassDef, _ast.Module)):
                    body = getattr(node, "body", [])
                    if body and isinstance(body[0], _ast.Expr) and isinstance(body[0].value, _ast.Constant):
                        for ln in range(body[0].lineno, body[0].end_lineno + 1):
                            doc_lines.add(ln)

            lines = source.splitlines()
            code_lines = []
            for i, line in enumerate(lines, 1):
                if i in doc_lines:
                    continue
                stripped = line.strip()
                if stripped.startswith("#"):
                    continue
                code_lines.append(line)
            code_text = "\n".join(code_lines)
            # 不得有 TB 写入模式
            tb_write = re.search(
                r"(INSERT\s+INTO\s+trial_balance|UPDATE\s+trial_balance|publish_to_tb|publishToTb)",
                code_text, re.I,
            )
            assert tb_write is None, (
                f"{aid}: provider 含 TB 写入逻辑 {tb_write.group()!r}（FC-9 红线）"
            )

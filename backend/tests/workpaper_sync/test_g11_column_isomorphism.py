# -*- coding: utf-8 -*-
"""G11-2 列同构判据 —— 单级表头 + 21 行 + 占比列引合计行 + 中性化后无公式的两列。

spec: `g-cycle-single-region-detail-lanes` · Task 11 / C-10
　　　设计依据 `evidence/task8-c6-remaining-eight-template-logic.md` §5 + 本轮逐格实测

═══ 五层闭环 ═══
  ① `header_text` 逐列 == **模板 R9 逐格实测**（🔴 九条唯一的单级表头）
  ② `json_key` 集合/顺序 == 前端 `G11DetailRow` 的 13 个受管字段（双向相等）
  ③ 🔴 **中性化实测**：`G`/`K` 两列公式被摘 ⇒ 判 `auto_source` 而非 `formula`
  ④ 🔴 占比列引合计行的绝对分母（裁决 G1R-H5）+ footer shared **成员格**
  ⑤ 单区几何 / `R32` 手填行不受管 / `seq` 是真列 / 行身份 `id`

🔴 本文件里凡「中性化」相关断言一律**真跑** `neutralize_oo_crash_if_formulas`
（模板副本，就地改）—— 这条链是 G11 与前四条 lane 的根本差异，推演必错。
"""
from __future__ import annotations

import os
import re
import shutil
import sys
import tempfile
import zipfile
from pathlib import Path

import openpyxl
import pytest

_REPO = Path(__file__).resolve().parents[3]
_BACKEND = _REPO / "backend"
if str(_BACKEND) not in sys.path:
    sys.path.insert(0, str(_BACKEND))
os.environ.setdefault("DB_DISABLE_SSL", "True")

from app.services.workpaper_sync.phase5_g11_02_detail import (  # noqa: E402
    AUTO_SOURCE_COLUMNS_G1102,
    FIELD_SPECS_G1102,
    FIRST_DATA_ROW_G1102,
    FOOTER_ROW_G1102,
    FORMULA_COLUMNS_G1102,
    FORMULA_TEMPLATES_G1102,
    FRONTEND_ONLY_FIELDS_G1102,
    HEADER_ROW_G1102,
    LAST_DATA_ROW_G1102,
    MANAGED_SHEET_G1102,
    NEUTRALIZED_COLUMNS_G1102,
    PROFIT_TOTAL_ROW_G1102,
    RATIO_DENOMINATOR_REFS_G1102,
    SPEC_G1102,
    STORE_ITEM_ID_G1102,
)
from app.services.workpaper_sync.phase5_row_table_sheet import (  # noqa: E402
    iter_store_rows,
    merge_projection_into_store_rows,
)

TPL = _BACKEND / "wp_templates" / "G" / "G11 投资收益.xlsx"
#: 🔴 受管 sheet 的包内 XML 路径（按 workbook 顺序第 7 张）—— 中性化层面的实测要读原始 XML，
#:   因为 openpyxl **会自动展开 shared formula**，看不到成员格的自闭合形态。
SHEET_XML_G1102 = "xl/worksheets/sheet7.xml"

_FE = _REPO / "audit-platform/frontend/src/components/workpaper/composables"
FE_COMPOSABLE = _FE / "useG11DetailAnalysis.ts"

#: 前端行接口里**不受管**的 8 个字段（行身份 + 勾稽键 + 分组 + 阈值派生 + 骨架标记）
_NON_MANAGED_FE_FIELDS = set(FRONTEND_ONLY_FIELDS_G1102)

_DATA_ROWS = tuple(range(FIRST_DATA_ROW_G1102, LAST_DATA_ROW_G1102 + 1))
_CELL_RE = re.compile(
    r'<c r="([A-Z]+\d+)"(?P<attrs>[^>]*?)(?:/>|>(?P<inner>.*?)</c>)', re.S
)


@pytest.fixture(scope="module")
def wb():
    book = openpyxl.load_workbook(TPL, data_only=False)
    try:
        yield book
    finally:
        book.close()


@pytest.fixture(scope="module")
def ws(wb):
    return wb[MANAGED_SHEET_G1102]


def _cell_map(path: Path) -> dict[str, str]:
    """`ref -> <c> 内容`。自闭合格记空串（🔴 不这样做，正则会跨格吃到下一个公式）。"""
    with zipfile.ZipFile(path) as zf:
        xml = zf.read(SHEET_XML_G1102).decode("utf-8")
    return {m.group(1): (m.group("inner") or "") for m in _CELL_RE.finditer(xml)}


def _formula_counts(cells: dict[str, str]) -> dict[str, int]:
    from openpyxl.utils import get_column_letter

    return {
        get_column_letter(c): sum(
            1 for r in _DATA_ROWS if "<f" in cells.get(f"{get_column_letter(c)}{r}", "")
        )
        for c in range(1, 14)
    }


@pytest.fixture(scope="module")
def neutralized() -> tuple[dict[str, str], dict[str, str], tuple[str, ...]]:
    """真跑中性化：返回 `(before_cells, after_cells, touched_refs)`。"""
    from app.services.workpaper_sync.g7_oo_crash_if_neutralize import (
        neutralize_oo_crash_if_formulas,
    )

    with tempfile.TemporaryDirectory() as td:
        cp = Path(td) / "g11.xlsx"
        shutil.copy2(TPL, cp)
        before = _cell_map(cp)
        touched = neutralize_oo_crash_if_formulas(cp)
        after = _cell_map(cp)
    return before, after, touched


def _frontend_row_fields() -> list[str]:
    src = FE_COMPOSABLE.read_text(encoding="utf-8")
    m = re.search(r"export interface G11DetailRow \{(?P<body>.*?)\n\}", src, re.S)
    assert m, "找不到 G11DetailRow 接口 —— 前端结构变了，判据须改写"
    return [
        fm.group(1)
        for line in m.group("body").splitlines()
        if (fm := re.match(r"\s*(\w+)\??:", line))
    ]


# ════════════════════════════════════════════════════════════════════════════
# ① header_text 逐列 == 模板 R9 实测（🔴 单级表头）
# ════════════════════════════════════════════════════════════════════════════
class TestSingleLevelHeaderMatchesTemplateCellByCell:
    def test_13_columns_in_excel_order_a_to_m(self) -> None:
        from openpyxl.utils import get_column_letter

        cols = [f[1] for f in FIELD_SPECS_G1102]
        assert cols == [get_column_letter(i) for i in range(1, 14)], (
            f"列序必须是 A..M 连续 13 列，实得 {cols}"
        )

    def test_header_is_single_level_with_zero_merges_on_r8_r9(self, ws) -> None:
        """🔴 九条唯一的单级表头：R8/R9 **零合并格** ⇒ 没有任何横向分组。

        R8 是段标题「二、审计过程」（只占 A8，不是表头行）。别家（G9/G10/G8/G14）都有
        R9 横向合并的组标题 ⇒ 这里若出现合并格，说明模板改成两级了，`header_row` 单值
        的声明就必须拆成 `header_group_row`/`header_leaf_row`。
        """
        merged = sorted(
            r.coord for r in ws.merged_cells.ranges if r.min_row in (8, 9)
        )
        assert merged == [], f"R8/R9 出现合并格 {merged} ⇒ 单级表头前提被推翻"
        assert str(ws["A8"].value or "") == "二、审计过程"
        assert ws["B8"].value is None, "A8 的段标题不该横跨到 B8"

    def test_spec_declares_header_row_not_group_leaf_pair(self) -> None:
        assert SPEC_G1102.header_row == HEADER_ROW_G1102 == 9
        assert SPEC_G1102.header_group_row is None, "单级表头不得声明 group 行"
        assert SPEC_G1102.header_leaf_row is None, "单级表头不得声明 leaf 行"

    def test_effective_columns_equal_max_column_so_uuid_is_n(self, ws) -> None:
        """有效列 13 **恰等于** `max_column` ⇒ 两个 uuid 口径重合，取 N。"""
        from openpyxl.utils import get_column_letter

        effective = [
            get_column_letter(c)
            for c in range(1, ws.max_column + 1)
            if any(
                ws.cell(row=r, column=c).value not in (None, "")
                for r in (HEADER_ROW_G1102, *_DATA_ROWS, FOOTER_ROW_G1102)
            )
        ]
        assert effective == [get_column_letter(i) for i in range(1, 14)]
        assert ws.max_column == 13, "模板 max_column 变了 ⇒ uuid 列取值依据要重算"
        assert SPEC_G1102.uuid_col == "N"

    def test_template_has_zero_defined_names(self, wb) -> None:
        """G 循环 13 册全 0 definedName（GC 基线）—— 保持为无。"""
        assert len(wb.defined_names) == 0

    @pytest.mark.parametrize(
        "field", FIELD_SPECS_G1102, ids=[f[1] for f in FIELD_SPECS_G1102]
    )
    def test_header_text_equals_template_r9_cell(self, field, ws) -> None:
        from openpyxl.utils import column_index_from_string

        _key, col, _mode, _vt, _json_key, header_text, _group = field
        cell = ws.cell(row=HEADER_ROW_G1102, column=column_index_from_string(col))
        expect = str(cell.value or "").strip()
        assert header_text == expect, (
            f"{col}9 header_text 实得 {header_text!r}，模板逐格为 {expect!r}"
        )

    def test_no_field_declares_a_group_header_cell(self) -> None:
        """单级表头 ⇒ `group_header_cell` 必须**全空**（别家是组标题格坐标）。"""
        assert {f[6] for f in FIELD_SPECS_G1102} == {""}

    def test_two_ratio_columns_share_the_same_header_text(self, ws) -> None:
        """🔴 `G`/`K` 表头同为「各项目占比」⇒ 判据不能靠 header_text 区分这两列。"""
        assert ws["G9"].value == ws["K9"].value == "各项目占比"
        by_col = {f[1]: f for f in FIELD_SPECS_G1102}
        assert by_col["G"][5] == by_col["K"][5] == "各项目占比"
        assert by_col["G"][4] != by_col["K"][4], "json_key 必须能区分两列"


# ════════════════════════════════════════════════════════════════════════════
# ② json_key == 前端 G11DetailRow 的 13 个受管字段
# ════════════════════════════════════════════════════════════════════════════
class TestJsonKeysMatchFrontendRowInterface:
    def test_declared_json_keys_equal_frontend_managed_fields(self) -> None:
        declared = {f[4] for f in FIELD_SPECS_G1102}
        frontend = set(_frontend_row_fields()) - _NON_MANAGED_FE_FIELDS
        assert declared == frontend, (
            "声明的 json_key 与前端 G11DetailRow 受管字段必须**双向相等**。\n"
            f"  前端有而声明缺：{sorted(frontend - declared)}\n"
            f"  声明有而前端缺：{sorted(declared - frontend)}"
        )
        assert len(declared) == 13

    def test_frontend_field_order_matches_excel_column_order(self) -> None:
        """🔴 G11 是九条里**唯一前端零改动**的一家：13 个受管字段的声明顺序已与模板
        列序 A..M 逐列对齐（实测，不是改出来的）⇒ 本判据是「保持对齐」的守卫。
        """
        fe = [f for f in _frontend_row_fields() if f not in _NON_MANAGED_FE_FIELDS]
        assert fe == [f[4] for f in FIELD_SPECS_G1102], (
            "前端字段顺序应与模板列序一致。\n"
            f"  前端：{fe}\n  声明：{[f[4] for f in FIELD_SPECS_G1102]}"
        )

    def test_eight_non_template_fields_are_declared_and_present(self) -> None:
        """8 个非模板列字段：必须在前端接口里、且**不**在受管列里。"""
        declared = {f[4] for f in FIELD_SPECS_G1102}
        fe = set(_frontend_row_fields())
        assert len(FRONTEND_ONLY_FIELDS_G1102) == 8
        for f in FRONTEND_ONLY_FIELDS_G1102:
            assert f in fe, f"{f} 不在前端接口里 ⇒ 登记表过期"
            assert f not in declared, f"{f} 不是模板列，不应受管"

    def test_change_rate_is_frontend_only_because_template_has_only_amount(self) -> None:
        """模板只有变动额 `L`，没有变动率列 ⇒ `changeRate` 只能留在前端。"""
        assert "changeRate" in FRONTEND_ONLY_FIELDS_G1102
        assert "changeAmount" in {f[4] for f in FIELD_SPECS_G1102}
        headers = {f[5] for f in FIELD_SPECS_G1102}
        assert "变动额" in headers
        assert not any("变动率" in h or "变动比" in h for h in headers)

    def test_seq_is_a_real_template_column_unlike_g9_g10_g8(self, ws) -> None:
        """🔴 `seq` 在 G11 是**真列**（A 列「序号」，模板 A10-A30 预填 1..21）。

        G9/G10/G8 的同名字段是纯显示序号、不占模板列 —— 照抄那三家把 `seq` 排除在
        受管面外，会让 A 列脱管、OO 侧改序号不回流。
        """
        by_col = {f[1]: f for f in FIELD_SPECS_G1102}
        assert by_col["A"][4] == "seq"
        assert by_col["A"][3] == "integer"
        assert by_col["A"][2] == "editable"
        seeded = [ws.cell(row=r, column=1).value for r in _DATA_ROWS]
        assert seeded == list(range(1, 22)), f"A 列预填实得 {seeded}"

    def test_row_identity_field_id_is_not_a_managed_column(self) -> None:
        assert SPEC_G1102.row_identity_key == "id" == FRONTEND_ONLY_FIELDS_G1102[0]
        assert "id" not in {f[4] for f in FIELD_SPECS_G1102}

    def test_row_id_generator_suffix_length_is_four(self) -> None:
        """🔴 `generateId()` 取 `slice(2, 6)`（4 位）—— G9/G10/G8 是 3 位。

        随机后缀长度决定同毫秒内新增多行的碰撞概率；这里钉住实测值，改短要有依据。
        """
        src = FE_COMPOSABLE.read_text(encoding="utf-8")
        m = re.search(
            r"function generateId\(\): string \{\s*return `g11d-\$\{Date\.now\(\)"
            r"\.toString\(36\)\}\$\{Math\.random\(\)\.toString\(36\)\.slice\((\d+), (\d+)\)\}`",
            src,
        )
        assert m, "generateId 形态变了 —— 行身份生成方式要重新核"
        assert (int(m.group(1)), int(m.group(2))) == (2, 6)


# ════════════════════════════════════════════════════════════════════════════
# ③ 🔴 中性化实测 ⇒ G/K 判 auto_source（本 lane 与前四条的根本差异）
# ════════════════════════════════════════════════════════════════════════════
class TestNeutralizationDecidesTheColumnModes:
    def test_before_neutralization_five_columns_carry_formulas(self, neutralized) -> None:
        """中性化**前**：`F`/`G`/`J`/`K`/`L` 各 21 格公式，`D`/`E`/`H`/`I` 是空可填格。"""
        before, _after, _touched = neutralized
        counts = _formula_counts(before)
        assert {c: n for c, n in counts.items() if n} == {
            "F": 21,
            "G": 21,
            "J": 21,
            "K": 21,
            "L": 21,
        }, f"数据区公式分布实得 {counts}"

    def test_after_neutralization_only_g_and_k_lose_their_formulas(
        self, neutralized
    ) -> None:
        """🔴 中性化**后**：`G`/`K` 归零，`F`/`J`/`L` 保持 21 —— 这就是两族声明的根据。"""
        _before, after, _touched = neutralized
        counts = _formula_counts(after)
        assert {c: n for c, n in counts.items() if n} == {"F": 21, "J": 21, "L": 21}
        for col in NEUTRALIZED_COLUMNS_G1102:
            assert counts[col] == 0, f"{col} 列中性化后仍有 {counts[col]} 格公式"

    def test_neutralizer_touches_exactly_44_cells_on_the_managed_sheet(
        self, neutralized
    ) -> None:
        """受管 sheet 恰 44 格 = 数据区 21×2 + footer `G31`/`K31` 两格成员格。"""
        _before, _after, touched = neutralized
        mine = [t for t in touched if SHEET_XML_G1102.split("/")[-1] in t]
        assert len(mine) == 44, f"受管 sheet 触及 {len(mine)} 格：{mine[:6]}…"
        assert {t.rsplit("!", 1)[1][0] for t in mine} == {"G", "K"}
        assert len(touched) > len(mine), "整册还有别的 sheet 带裸 IF（G11-4/G11-1）"

    def test_managed_sheet_is_the_only_g_cycle_lane_hit_by_neutralization(
        self, neutralized
    ) -> None:
        """🔴 与 G9/G10/G8/G14 的根本差异：那四家的受管表零裸 IF，中性化只动审定表。"""
        _before, _after, touched = neutralized
        assert any(SHEET_XML_G1102.split("/")[-1] in t for t in touched), (
            "受管表不再被中性化命中 ⇒ G/K 应改判 formula，本文件多条判据要翻"
        )

    def test_ratio_columns_are_declared_auto_source_not_formula(self) -> None:
        """🔴 `G`/`K` 判 `auto_source`：中性化后没有公式，判 `formula` 会抛。

        `excel_materialize` 两条检查方向相反 ——
          * `mode=formula`     ⇒ 要求 `view.has_formula` 为真
          * `mode=auto_source` ⇒ 要求该格**不是**公式
        中性化把 `G`/`K` 的 `<f>` 摘干净了 ⇒ 只有 `auto_source` 能过。
        """
        by_col = {f[1]: f for f in FIELD_SPECS_G1102}
        for col in ("G", "K"):
            assert by_col[col][2] == "auto_source", (
                f"{col} 列 mode 实得 {by_col[col][2]!r} —— 判 formula 会在 materialize "
                f"抛 ProtectedRegionWriteError（substrate 无公式）"
            )
        assert AUTO_SOURCE_COLUMNS_G1102 == ("G", "K") == NEUTRALIZED_COLUMNS_G1102

    def test_auto_source_is_protected_so_oo_edits_do_not_merge(self) -> None:
        """`auto_source` 与 `formula` 同属 `PROTECTED_MODES` ⇒ 保护力度**没有**放宽。"""
        from app.services.workpaper_sync.contracts import PROTECTED_MODES, FieldMode

        assert FieldMode.auto_source in PROTECTED_MODES
        assert FieldMode.formula in PROTECTED_MODES
        assert FieldMode.editable not in PROTECTED_MODES

    def test_formula_columns_exclude_the_two_ratio_columns(self) -> None:
        assert FORMULA_COLUMNS_G1102 == ("F", "J", "L")
        assert not set(FORMULA_COLUMNS_G1102) & set(AUTO_SOURCE_COLUMNS_G1102)
        assert sorted(FORMULA_TEMPLATES_G1102) == ["F", "J", "L"], (
            "占比列不进 formula_templates —— 那张表驱动 mode=formula 的逐格比对"
        )

    def test_formula_mask_covers_exactly_the_three_formula_columns(self) -> None:
        """`formula_mask` 由 `formula_columns` 派生 ⇒ `G`/`K` 不在 mask 内（正确）。"""
        assert SPEC_G1102.formula_mask == (
            "F10:F30",
            "J10:J30",
            "L10:L30",
        ), f"实得 {SPEC_G1102.formula_mask}"

    def test_every_formula_mode_field_is_covered_by_formula_mask(self) -> None:
        """契约层 CS-13 的本地复核（提前于 parse_contract 打红，定位更快）。"""
        from app.services.workpaper_sync.contracts import column_in_ranges

        mask = SPEC_G1102.formula_mask
        for key, col, mode, _vt, _json, _hdr, _grp in FIELD_SPECS_G1102:
            if mode != "formula":
                continue
            assert column_in_ranges(col, mask), (
                f"{key}（列 {col}）声明 mode=formula 但不在 formula_mask {mask} 内"
            )

    @pytest.mark.parametrize("col", sorted(FORMULA_TEMPLATES_G1102))
    def test_formula_template_renders_to_every_data_row(self, col: str, ws) -> None:
        """三列公式逐格可比（中性化前的模板原式）。"""
        from openpyxl.utils import column_index_from_string

        ci = column_index_from_string(col)
        for r in _DATA_ROWS:
            got = ws.cell(row=r, column=ci).value
            assert got == FORMULA_TEMPLATES_G1102[col].format(r=r), (
                f"{col}{r} 模板实得 {got!r}，声明渲染为 "
                f"{FORMULA_TEMPLATES_G1102[col].format(r=r)!r}"
            )

    def test_the_three_formula_columns_carry_no_bare_if(self) -> None:
        """`F`/`J`/`L` 不含 `IF(` —— 这是它们能留在 `formula` 族的唯一理由。"""
        for col, tpl in FORMULA_TEMPLATES_G1102.items():
            assert not re.search(r"\bIF\(", tpl), f"{col} 含裸 IF ⇒ 会被中性化摘掉"


# ════════════════════════════════════════════════════════════════════════════
# ④ 🔴 占比列引合计行的绝对分母（G1R-H5）+ footer shared 成员格
# ════════════════════════════════════════════════════════════════════════════
class TestRatioColumnsReferenceTheFooterRow:
    @pytest.mark.parametrize("col", sorted(RATIO_DENOMINATOR_REFS_G1102))
    def test_master_cell_formula_is_recorded_verbatim(self, col: str, ws) -> None:
        """🔴 逐格比对 shared 组**主格**（`G10`/`K10`）—— 看成员格会得到展开式而误判。"""
        verbatim, _denom = RATIO_DENOMINATOR_REFS_G1102[col]
        got = ws[f"{col}{FIRST_DATA_ROW_G1102}"].value
        assert got == f"={verbatim}", f"{col}10 实得 {got!r}，声明为 {'=' + verbatim!r}"

    def test_the_two_denominators_are_different_columns(self) -> None:
        """🔴 spec 原文只说「引 F31」；Task 2 实测 `K` 引 `$J$31` ⇒ 两个分母**不同**。"""
        denoms = {c: d for c, (_v, d) in RATIO_DENOMINATOR_REFS_G1102.items()}
        assert denoms == {"G": "$F$31", "K": "$J$31"}
        assert len(set(denoms.values())) == 2, "两列分母相同 ⇒ 说明有一列写错了"

    @pytest.mark.parametrize("col", sorted(RATIO_DENOMINATOR_REFS_G1102))
    def test_denominator_points_at_the_real_footer_row(self, col: str) -> None:
        """分母行号必须**恰是** footer 行 —— 引数据区任何一行都是错的。"""
        _verbatim, denom = RATIO_DENOMINATOR_REFS_G1102[col]
        m = re.fullmatch(r"\$([A-Z]+)\$(\d+)", denom)
        assert m, f"{col} 的分母 {denom!r} 不是绝对引用"
        assert int(m.group(2)) == FOOTER_ROW_G1102 == 31
        assert m.group(1) in FORMULA_COLUMNS_G1102, (
            f"{col} 的分母列 {m.group(1)} 不是审定数列 ⇒ 占比口径错"
        )

    @pytest.mark.parametrize("col", sorted(RATIO_DENOMINATOR_REFS_G1102))
    def test_denominator_stays_absolute_across_every_data_row(
        self, col: str, ws
    ) -> None:
        """数据区 21 行的分子随行走、分母**恒定**（shared 组的正确展开形态）。"""
        _verbatim, denom = RATIO_DENOMINATOR_REFS_G1102[col]
        num_col = denom[1]  # $F$31 -> F
        for r in _DATA_ROWS:
            got = str(ws[f"{col}{r}"].value or "")
            assert got == f"=IF({num_col}{r}=0,0,{num_col}{r}/{denom})", (
                f"{col}{r} 实得 {got!r}"
            )

    def test_footer_ratio_cells_are_self_closing_shared_members(
        self, neutralized
    ) -> None:
        """🔴 `G31`/`K31` 是自闭合 `<f t="shared" si="N"/>`，**不带公式文本**。

        openpyxl 会自动展开成 `=IF(F31=0,0,F31/$F$31)` —— 判「占比列是否指向真实合计行」
        必须看 XML 的 shared 组主格；看成员格的 openpyxl 值会被展开式骗过去
        （Task 5 首版判据踩过这个坑）。
        """
        before, _after, _touched = neutralized
        for col in ("G", "K"):
            inner = before[f"{col}{FOOTER_ROW_G1102}"]
            assert re.search(r'<f t="shared" si="\d+"/>', inner), (
                f"{col}31 实得 {inner[:70]!r} —— 不再是成员格"
            )

    def test_openpyxl_expands_the_member_cell_so_xml_is_the_authority(self, ws) -> None:
        """把上一条的「为什么必须读 XML」钉成可执行事实（防后人改回 openpyxl 口径）。"""
        assert ws[f"G{FOOTER_ROW_G1102}"].value == "=IF(F31=0,0,F31/$F$31)"
        assert ws[f"K{FOOTER_ROW_G1102}"].value == "=IF(J31=0,0,J31/$J$31)"

    def test_footer_sum_columns_are_plain_formulas_not_shared(
        self, neutralized
    ) -> None:
        """footer 的 7 个 `SUM` 列是普通 `<f>`，中性化**不动**它们。"""
        before, after, _touched = neutralized
        for col in ("D", "E", "F", "H", "I", "J", "L"):
            body = f"<f>SUM({col}{FIRST_DATA_ROW_G1102}:{col}{LAST_DATA_ROW_G1102})</f>"
            assert body in before[f"{col}{FOOTER_ROW_G1102}"]
            assert body in after[f"{col}{FOOTER_ROW_G1102}"], (
                f"{col}31 的 SUM 被中性化摘了 ⇒ footer 合计会失效"
            )

    def test_footer_carries_total_formula_is_declared(self, ws) -> None:
        """`footer_carries_total_formula=True` ⇒ `_grow_managed_table_ref` 归一化区间。"""
        assert SPEC_G1102.footer_carries_total_formula is True
        assert str(ws[f"A{FOOTER_ROW_G1102}"].value) == SPEC_G1102.footer_marker == "合计"


# ════════════════════════════════════════════════════════════════════════════
# ⑤ 几何 / R32 手填行 / store 读写
# ════════════════════════════════════════════════════════════════════════════
class TestGeometryAndStore:
    def test_single_region_r10_to_r30_is_21_rows(self) -> None:
        """九条里最长的数据区 —— 21 行。"""
        assert (SPEC_G1102.first_data_row, SPEC_G1102.last_data_row) == (10, 30)
        assert len(_DATA_ROWS) == 21
        assert SPEC_G1102.footer_row == FOOTER_ROW_G1102 == 31
        assert SPEC_G1102.footer_row == SPEC_G1102.last_data_row + 1
        assert SPEC_G1102.store_item_id == STORE_ITEM_ID_G1102 == "G11-detail-rows"
        assert not SPEC_G1102.row_section_field, "单区不得声明 row_section_field"

    def test_profit_total_row_is_hand_filled_and_out_of_scope(self, ws) -> None:
        """🔴 `R32`「本年利润总额」是手填分析行：既不是 footer、也不受管。

        它有**独立 store 键** `G11-detail-profit-total`（前端 `ITEM_ID_PROFIT`）⇒ 受管区
        必须止于 R30、footer 必须是 R31，否则行表引擎会把 R32 当数据行吞掉。
        """
        assert PROFIT_TOTAL_ROW_G1102 == 32 == FOOTER_ROW_G1102 + 1
        assert str(ws[f"A{PROFIT_TOTAL_ROW_G1102}"].value) == "本年利润总额"
        assert PROFIT_TOTAL_ROW_G1102 > SPEC_G1102.footer_row
        src = FE_COMPOSABLE.read_text(encoding="utf-8")
        assert "const ITEM_ID_PROFIT = 'G11-detail-profit-total'" in src
        assert f"const ITEM_ID_ROWS = '{STORE_ITEM_ID_G1102}'" in src

    def test_audit_note_rows_below_are_out_of_scope(self, ws) -> None:
        assert str(ws["A33"].value or "").startswith("三、审计说明")
        assert 33 > PROFIT_TOTAL_ROW_G1102

    def test_row_identity_key_is_id_not_rowkey(self) -> None:
        """🔴 行身份是 `id`；同接口里的 `rowKey` 是与 G11-1 对齐的**勾稽键**，不是身份。"""
        assert SPEC_G1102.row_identity_key == "id"
        assert "rowKey" in FRONTEND_ONLY_FIELDS_G1102

    def test_iter_reads_every_row_by_id(self) -> None:
        payload = [{"id": "g11d-a"}, {"id": "g11d-b"}, {"id": "g11d-sk-equity"}]
        assert [rid for rid, _ in iter_store_rows(SPEC_G1102, payload)] == [
            "g11d-a",
            "g11d-b",
            "g11d-sk-equity",
        ]

    def test_duplicate_id_fails_closed(self) -> None:
        from app.services.workpaper_sync.phase5_row_table_sheet import (
            RowTableStorePayloadError,
        )

        with pytest.raises(RowTableStorePayloadError, match="重复行身份"):
            list(iter_store_rows(SPEC_G1102, [{"id": "x"}, {"id": "x"}]))

    def test_merge_keys_rows_by_id_and_does_not_touch_ratio_columns(self) -> None:
        """回流只合并 `editable` 字段；`G`/`K` 受保护 ⇒ 不进 store。"""

        class _FV:
            def __init__(self, key: str, row_key: str, value):
                self.stable_key = key
                self.row_key = row_key
                self.value = value
                self.is_protected = False

        key = f"{SPEC_G1102.table_key}/g11d-a/current_unadjusted"

        class _Proj:
            def stable_keys(self):
                return [key]

            def get(self, k):
                return _FV(k, "g11d-a", 8888) if k == key else None

        rows, applied, _visited, touched = merge_projection_into_store_rows(
            SPEC_G1102,
            projection=_Proj(),
            base_rows=[{"id": "g11d-a"}, {"id": "g11d-b"}],
        )
        assert applied == 1 and touched == {"g11d-a"}
        hit = next(r for r in rows if r["id"] == "g11d-a")
        assert hit["currentUnadjusted"] == 8888
        assert "currentShare" not in hit, "占比列受保护，不该出现在 store 行里"
        assert [r["id"] for r in rows] == ["g11d-a", "g11d-b"]

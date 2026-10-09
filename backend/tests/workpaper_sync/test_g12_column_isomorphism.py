# -*- coding: utf-8 -*-
"""G12-2 列同构判据 —— 两级表头 + 零公式列 + 五处模板缺陷 + 打断的 footer shared 组。

spec: `g-cycle-single-region-detail-lanes` · Task 13 / C-12

═══ 五层闭环 ═══
  ① `header_text` 逐列 == 模板两级表头实测（🔴 横向组只有 `D7:G7`）+ 🔴 有效列 < max_column
  ② `json_key` 集合/顺序 == 前端 `G12HedgeDetailRow` 的 10 个受管字段
  ③ 🔴 `formula_columns` **为空**：`G`/`I` 判 `editable` 的裁决取证
  ④ 🔴 五处模板缺陷台账逐格比对 + footer `F14:J14` 组被 `H14`/`I14` 打断
  ⑤ 几何（R9-R13，只 R9/R10 预填）/ 零裸 IF / store 读写
"""
from __future__ import annotations

import os
import re
import shutil
import sys
import tempfile
import zipfile
from pathlib import Path

import openpyxl
import pytest

_REPO = Path(__file__).resolve().parents[3]
_BACKEND = _REPO / "backend"
if str(_BACKEND) not in sys.path:
    sys.path.insert(0, str(_BACKEND))
os.environ.setdefault("DB_DISABLE_SSL", "True")

from app.services.workpaper_sync.phase5_g12_02_detail import (  # noqa: E402
    BLANK_DATA_ROWS_G1202,
    BOOLEAN_COLUMNS_G1202,
    FIELD_SPECS_G1202,
    FIRST_DATA_ROW_G1202,
    FOOTER_ROW_G1202,
    FOOTER_SHAPE_FACTS_G1202,
    FORMULA_COLUMNS_G1202,
    FRONTEND_DERIVED_COLUMNS_G1202,
    FRONTEND_ONLY_FIELDS_G1202,
    LAST_DATA_ROW_G1202,
    MANAGED_SHEET_G1202,
    SEEDED_DATA_ROWS_G1202,
    SPEC_G1202,
    STORE_ITEM_ID_G1202,
    TEMPLATE_FOOTER_FORMULAS_G1202,
    TEMPLATE_FORMULA_COVERAGE_DEFECTS_G1202,
    TEMPLATE_ROW_FORMULAS_G1202,
)
from app.services.workpaper_sync.phase5_row_table_sheet import (  # noqa: E402
    iter_store_rows,
    merge_projection_into_store_rows,
)

TPL = _BACKEND / "wp_templates" / "G" / "G12 净敞口套期收益.xlsx"
SHEET_XML_G1202 = "xl/worksheets/sheet7.xml"

_FE = _REPO / "audit-platform/frontend/src/components/workpaper/composables"
FE_COMPOSABLE = _FE / "useG12HedgeDetail.ts"
FE_CALC = _FE / "g12NetHedgeDetailCalc.ts"

_NON_MANAGED_FE_FIELDS = set(FRONTEND_ONLY_FIELDS_G1202)
_DATA_ROWS = tuple(range(FIRST_DATA_ROW_G1202, LAST_DATA_ROW_G1202 + 1))
_CELL_RE = re.compile(
    r'<c r="([A-Z]+\d+)"(?P<attrs>[^>]*?)(?:/>|>(?P<inner>.*?)</c>)', re.S
)


@pytest.fixture(scope="module")
def wb():
    book = openpyxl.load_workbook(TPL, data_only=False)
    try:
        yield book
    finally:
        book.close()


@pytest.fixture(scope="module")
def ws(wb):
    return wb[MANAGED_SHEET_G1202]


def _cell_map(path: Path) -> dict[str, str]:
    with zipfile.ZipFile(path) as zf:
        xml = zf.read(SHEET_XML_G1202).decode("utf-8")
    return {m.group(1): (m.group("inner") or "") for m in _CELL_RE.finditer(xml)}


@pytest.fixture(scope="module")
def cells() -> dict[str, str]:
    return _cell_map(TPL)


@pytest.fixture(scope="module")
def neutralized() -> tuple[dict[str, str], tuple[str, ...]]:
    from app.services.workpaper_sync.g7_oo_crash_if_neutralize import (
        neutralize_oo_crash_if_formulas,
    )

    with tempfile.TemporaryDirectory() as td:
        cp = Path(td) / "g12.xlsx"
        shutil.copy2(TPL, cp)
        touched = neutralize_oo_crash_if_formulas(cp)
        after = _cell_map(cp)
    return after, touched


def _frontend_row_fields() -> list[str]:
    src = FE_COMPOSABLE.read_text(encoding="utf-8")
    m = re.search(r"export interface G12HedgeDetailRow \{(?P<body>.*?)\n\}", src, re.S)
    assert m, "找不到 G12HedgeDetailRow 接口 —— 前端结构变了，判据须改写"
    return [
        fm.group(1)
        for line in m.group("body").splitlines()
        if (fm := re.match(r"\s*(\w+)\??:", line))
    ]


# ════════════════════════════════════════════════════════════════════════════
# ① 两级表头 + 有效列 < max_column
# ════════════════════════════════════════════════════════════════════════════
class TestTwoLevelHeaderAndTrailingBlankColumns:
    def test_10_columns_in_excel_order_a_to_j(self) -> None:
        from openpyxl.utils import get_column_letter

        cols = [f[1] for f in FIELD_SPECS_G1202]
        assert cols == [get_column_letter(i) for i in range(1, 11)], (
            f"列序必须是 A..J 连续 10 列，实得 {cols}"
        )

    def test_header_is_two_level_at_r7_r8_not_absent(self, ws) -> None:
        """🔴 spec 原文的「无表头行（R9 即数据）」前提**不成立**。

        裁决 G1R-H2（要框架层加 `has_header_row`）因此不适用，本 lane 走通例分支。
        判据钉住 R7/R8 真的是表头：R7 有组标题、R8 有叶子、R9 起才是数据。
        """
        assert SPEC_G1202.header_group_row == 7
        assert SPEC_G1202.header_leaf_row == 8
        assert SPEC_G1202.header_row is None
        assert SPEC_G1202.first_data_row == 9
        assert str(ws["A7"].value or "") == "项目"
        assert str(ws["G8"].value or "") == "校验"
        assert str(ws["A6"].value or "") == "二、审计过程："
        assert ws["B6"].value is None, "段标题不该横跨到 B6"

    def test_exactly_one_horizontal_group_on_r7(self, ws) -> None:
        """🔴 横向组**只有一个** `D7:G7`；其余六列是 R7:R8 纵向合并。"""
        groups = sorted(
            r.coord
            for r in ws.merged_cells.ranges
            if r.min_row == 7 and r.max_row == 7 and r.max_col > r.min_col
        )
        assert groups == ["D7:G7"], f"R7 横向组实得 {groups}"
        assert str(ws["D7"].value or "") == "套期工具公允价值"
        assert {f[6] for f in FIELD_SPECS_G1202 if f[6]} == {"D7"}

    def test_six_columns_are_vertically_merged_over_r7_r8(self, ws) -> None:
        vertical = sorted(
            r.coord
            for r in ws.merged_cells.ranges
            if r.min_row == 7 and r.max_row == 8 and r.min_col == r.max_col
        )
        assert vertical == ["A7:A8", "B7:B8", "C7:C8", "H7:H8", "I7:I8", "J7:J8"]
        by_col = {f[1]: f for f in FIELD_SPECS_G1202}
        for col in ("A", "B", "C", "H", "I", "J"):
            assert by_col[col][6] == "", f"{col} 是纵向合并，不该声明横向组标题格"

    @pytest.mark.parametrize(
        "field", FIELD_SPECS_G1202, ids=[f[1] for f in FIELD_SPECS_G1202]
    )
    def test_header_text_equals_template_cell(self, field, ws) -> None:
        from openpyxl.utils import column_index_from_string

        _key, col, _mode, _vt, _json_key, header_text, _group = field
        ci = column_index_from_string(col)
        leaf = ws.cell(row=8, column=ci).value
        group = ws.cell(row=7, column=ci).value
        expect = (
            str(leaf).strip()
            if (leaf and str(leaf).strip())
            else str(group or "").strip()
        )
        assert header_text == expect, (
            f"{col} 列 header_text 实得 {header_text!r}，模板逐格为 {expect!r}"
            f"（R7={group!r} / R8={leaf!r}）"
        )

    def test_effective_columns_are_fewer_than_max_column_so_uuid_is_k(self, ws) -> None:
        """🔴 有效列 **10**（A..J）< `max_column=15`（5 个空尾列）⇒ uuid 取 `K` 不是 `P`。

        取 `P` 会把 5 个空列圈进受管区。九条里第二家出现这种情况（另一家是 G8 的 23/24）。
        """
        from openpyxl.utils import get_column_letter

        effective = [
            get_column_letter(c)
            for c in range(1, ws.max_column + 1)
            if any(
                ws.cell(row=r, column=c).value not in (None, "")
                for r in (7, 8, *_DATA_ROWS, FOOTER_ROW_G1202)
            )
        ]
        assert effective == [get_column_letter(i) for i in range(1, 11)]
        assert ws.max_column == 15, "模板 max_column 变了 ⇒ uuid 列取值依据要重算"
        assert len(effective) < ws.max_column
        assert SPEC_G1202.uuid_col == "K"

    def test_template_has_zero_defined_names(self, wb) -> None:
        assert len(wb.defined_names) == 0

    def test_workbook_has_eleven_sheets(self, wb) -> None:
        """册内 sheet 数进判据 —— 灰度开关注释里逐字列了 11 张，改了要同步。"""
        assert len(wb.sheetnames) == 11
        assert MANAGED_SHEET_G1202 in wb.sheetnames


# ════════════════════════════════════════════════════════════════════════════
# ② json_key == 前端 G12HedgeDetailRow 的受管字段
# ════════════════════════════════════════════════════════════════════════════
class TestJsonKeysMatchFrontendRowInterface:
    def test_declared_json_keys_equal_frontend_managed_fields(self) -> None:
        declared = {f[4] for f in FIELD_SPECS_G1202}
        frontend = set(_frontend_row_fields()) - _NON_MANAGED_FE_FIELDS
        assert declared == frontend, (
            "声明的 json_key 与前端 G12HedgeDetailRow 受管字段必须**双向相等**。\n"
            f"  前端有而声明缺：{sorted(frontend - declared)}\n"
            f"  声明有而前端缺：{sorted(declared - frontend)}"
        )
        assert len(declared) == 10

    def test_frontend_field_order_matches_excel_column_order(self) -> None:
        fe = [f for f in _frontend_row_fields() if f not in _NON_MANAGED_FE_FIELDS]
        assert fe == [f[4] for f in FIELD_SPECS_G1202], (
            f"前端字段顺序应与模板列序一致。\n  前端：{fe}\n  声明：{[f[4] for f in FIELD_SPECS_G1202]}"
        )

    def test_four_non_template_fields_are_declared_and_present(self) -> None:
        declared = {f[4] for f in FIELD_SPECS_G1202}
        fe = set(_frontend_row_fields())
        assert len(FRONTEND_ONLY_FIELDS_G1202) == 4
        for f in FRONTEND_ONLY_FIELDS_G1202:
            assert f in fe, f"{f} 不在前端接口里 ⇒ 登记表过期"
            assert f not in declared, f"{f} 不是模板列，不应受管"

    def test_seq_is_display_only_unlike_g11(self, ws) -> None:
        """🔴 `seq` 是纯显示序号 —— 模板 A 列是「项目」不是序号（与 G11 的 seq 是真列相反）。"""
        assert "seq" in FRONTEND_ONLY_FIELDS_G1202
        assert str(ws["A7"].value or "") == "项目"
        by_col = {f[1]: f for f in FIELD_SPECS_G1202}
        assert by_col["A"][4] == "item"
        assert by_col["A"][3] == "text"

    def test_row_kind_drives_two_row_families_and_stays_unmanaged(self) -> None:
        """`rowKind` 决定两类行各自的取值口径（FV 分配 / 摊销），但模板没有这一列。"""
        assert "rowKind" in FRONTEND_ONLY_FIELDS_G1202
        src = FE_COMPOSABLE.read_text(encoding="utf-8")
        assert "rowKind === 'amortization' ? 'amortization' : 'fv_allocation'" in src
        calc = FE_CALC.read_text(encoding="utf-8")
        assert "export type G12NetHedgeDetailRowKind = 'fv_allocation' | 'amortization'" in calc

    def test_row_identity_generator_suffix_length_is_three(self) -> None:
        """🔴 后缀取 **3** 位（`slice(2, 5)`）—— G11/G13 取 4 位，形态判定必须按值 grep。"""
        src = FE_COMPOSABLE.read_text(encoding="utf-8")
        m = re.search(
            r"function genId\(\) \{\s*return `g12h-\$\{Date\.now\(\)\.toString\(36\)\}"
            r"\$\{Math\.random\(\)\.toString\(36\)\.slice\((\d+), (\d+)\)\}`",
            src,
        )
        assert m, "genId 形态变了 —— 行身份生成方式要重新核"
        assert (int(m.group(1)), int(m.group(2))) == (2, 5)


# ════════════════════════════════════════════════════════════════════════════
# ③ 🔴 零公式列：G/I 判 editable 的裁决取证
# ════════════════════════════════════════════════════════════════════════════
class TestZeroFormulaColumnsForcesEditable:
    def test_no_column_has_a_formula_on_every_data_row(self, cells) -> None:
        """🔴 裁决前提：数据区 R9-R13 **没有任何一列**每行都有公式。"""
        from openpyxl.utils import get_column_letter

        coverage = {
            get_column_letter(c): tuple(
                r for r in _DATA_ROWS
                if "<f" in cells.get(f"{get_column_letter(c)}{r}", "")
            )
            for c in range(1, 11)
        }
        nonempty = {k: v for k, v in coverage.items() if v}
        assert nonempty == {"G": (9,), "I": (9, 10)}, f"数据区公式分布实得 {nonempty}"
        assert not any(len(v) == len(_DATA_ROWS) for v in coverage.values())

    def test_formula_columns_is_empty(self) -> None:
        assert FORMULA_COLUMNS_G1202 == ()
        assert SPEC_G1202.formula_mask == ()
        assert SPEC_G1202.formula_templates == {}

    def test_all_ten_columns_are_editable(self) -> None:
        modes = {f[1]: f[2] for f in FIELD_SPECS_G1202}
        assert set(modes.values()) == {"editable"}, f"实得 {sorted(set(modes.values()))}"

    def test_g_and_i_are_editable_because_template_coverage_is_partial(self) -> None:
        """🔴 `G`/`I` 判 `editable`：模板公式只覆盖 1/5 与 2/5 行。

        判 `formula` 要求每格 `view.has_formula` 为真 ⇒ 会在 `G10:G13` 四格与 `I11:I13`
        三格抛 `ProtectedRegionWriteError`。
        """
        modes = {f[1]: f[2] for f in FIELD_SPECS_G1202}
        for col in FRONTEND_DERIVED_COLUMNS_G1202:
            assert modes[col] == "editable"
        assert FRONTEND_DERIVED_COLUMNS_G1202 == ("G", "I")

    def test_existing_formulas_are_plain_not_shared_masters(self, cells) -> None:
        """🔴 三格是**普通公式**而非 shared 主格 ⇒ 写字面量不撞 `SharedFormulaMasterWriteError`。"""
        for (col, row) in TEMPLATE_ROW_FORMULAS_G1202:
            inner = cells[f"{col}{row}"]
            assert "<f>" in inner, f"{col}{row} 不再是普通 <f> 形态：{inner[:60]!r}"
            assert 't="shared"' not in inner, (
                f"{col}{row} 变成了 shared formula ⇒ 判 editable 会撞 "
                "SharedFormulaMasterWriteError，裁决必须重做"
            )

    @pytest.mark.parametrize(
        "key", sorted(TEMPLATE_ROW_FORMULAS_G1202), ids=lambda k: f"{k[0]}{k[1]}"
    )
    def test_row_formula_ledger_matches_template_cell_by_cell(self, key, ws) -> None:
        col, row = key
        got = ws[f"{col}{row}"].value
        assert got == TEMPLATE_ROW_FORMULAS_G1202[key], (
            f"{col}{row} 模板实得 {got!r}，台账为 {TEMPLATE_ROW_FORMULAS_G1202[key]!r}"
        )

    def test_g_is_a_boolean_check_column_with_editable_mode(self) -> None:
        """🔴 裁决 G1R-H4 的第三个位点：`boolean` 成立、`formula` **不成立**。"""
        by_col = {f[1]: f for f in FIELD_SPECS_G1202}
        assert by_col["G"][3] == "boolean"
        assert by_col["G"][2] == "editable"
        assert BOOLEAN_COLUMNS_G1202 == ("G",)
        assert [f[1] for f in FIELD_SPECS_G1202 if f[3] == "boolean"] == ["G"]
        assert TEMPLATE_ROW_FORMULAS_G1202[("G", 9)] == "=D9=SUM(E9:F9)"

    def test_derived_columns_are_persisted_by_the_frontend(self) -> None:
        """🔴 `fvCheck`/`netHedgePnl` 必须真的落库 —— 否则 R11-R13 在 Excel 里永远空。"""
        src = FE_COMPOSABLE.read_text(encoding="utf-8")
        assert "fvCheck: boolean" in src and "netHedgePnl: number" in src
        # 🔴 只断言「两个字段各自出现在 persist 的 payload 里」，不写死整行字面量 ——
        #    写死会在字段重排时假红（本轮按模板列序把 fvCheck 挪到 hedgeAdjAmortization 前面，
        #    首版判据就是这么红的）。
        m = re.search(r"function persist\(\) \{(?P<body>[\s\S]*?)\n  \}", src)
        assert m, "找不到 persist()"
        body = m.group("body")
        for field in ("fvCheck: r.fvCheck", "netHedgePnl: r.netHedgePnl"):
            assert field in body, f"persist() 未落库 {field} ⇒ 受管面拿不到值"
        # 单一真源仍是那两个纯函数（不许在 enrich 里重写口径）
        assert "calcFvAllocationCheck(instrumentFvCumulative, salesPortion, purchasePortion)" in src
        assert "netHedgePnl: calcNetHedgePnl(calcInput)" in src
        # rowCalcs 改为从行模型读，对外字段名不变（消费方零改动）
        assert "fvCheckOk: r.fvCheck," in src

    def test_derived_column_semantics_match_the_template_formulas(self) -> None:
        """前端纯函数的口径必须与模板公式一致（`G=D==E+F` / `I=E+H`）。"""
        calc = FE_CALC.read_text(encoding="utf-8")
        assert "Math.abs(salesPortion + purchasePortion - instrumentFvCumulative) <= TOLERANCE" in calc
        assert TEMPLATE_ROW_FORMULAS_G1202[("I", 9)] == "=E9+H9"
        # I 列：FV 分配行取销售部分（E），摊销行取摊销额（H）—— 与 `=E+H` 在各自行上等价
        assert "if (row.rowKind === 'amortization') return parseNum(row.hedgeAdjAmortization)" in calc
        assert "return parseNum(row.salesPortion)" in calc


# ════════════════════════════════════════════════════════════════════════════
# ④ 🔴 五处模板缺陷 + 打断的 footer shared 组
# ════════════════════════════════════════════════════════════════════════════
class TestFiveTemplateDefectsAndInterruptedFooterGroup:
    def test_defect_ledger_has_exactly_five_entries(self) -> None:
        assert len(TEMPLATE_FORMULA_COVERAGE_DEFECTS_G1202) == 5
        located = [row[1] for row in TEMPLATE_FORMULA_COVERAGE_DEFECTS_G1202]
        assert located == ["B14", "I14", "G10:G13", "I11:I13", "H14"]
        # 每条都要写出「会计正确口径 / 应有形态」，否则台账只是抱怨
        for tag, where, what, want in TEMPLATE_FORMULA_COVERAGE_DEFECTS_G1202:
            assert tag and where and what and want, f"台账 {where} 有空字段"

    def test_defect_1_b14_misses_two_rows_and_is_a_value_error(self, ws) -> None:
        """🔴 缺陷①：`B14=SUM(B9,B12,B13:B13)` 漏加 B10/B11 —— 这是**数值错**不是形态问题。"""
        assert ws[f"B{FOOTER_ROW_G1202}"].value == "=SUM(B9,B12,B13:B13)"
        assert TEMPLATE_FOOTER_FORMULAS_G1202["B"] == "=SUM(B9,B12,B13:B13)"
        cited = {int(x) for x in re.findall(r"B(\d+)", "=SUM(B9,B12,B13:B13)")}
        assert cited == {9, 12, 13}
        assert {10, 11} & set(_DATA_ROWS) and not ({10, 11} & cited), "漏加的正是 B10/B11"
        entry = next(r for r in TEMPLATE_FORMULA_COVERAGE_DEFECTS_G1202 if r[1] == "B14")
        assert "数值错" in entry[2]
        assert entry[3] == "应为 =SUM(B9:B13)"

    def test_defect_2_i14_starts_above_the_data_region(self, ws) -> None:
        """🔴 缺陷②：`I14=SUM(I7:I13)` 起点 R7 是**表头组行**，不是数据行。"""
        assert ws[f"I{FOOTER_ROW_G1202}"].value == "=SUM(I7:I13)"
        assert TEMPLATE_FOOTER_FORMULAS_G1202["I"] == "=SUM(I7:I13)"
        start = int(re.search(r"SUM\(I(\d+):", "=SUM(I7:I13)").group(1))  # type: ignore[union-attr]
        assert start == SPEC_G1202.header_group_row == 7
        assert start < FIRST_DATA_ROW_G1202

    def test_defect_3_and_4_are_fill_down_gaps(self, cells) -> None:
        """🔴 缺陷③④：`G` 只填 R9、`I` 只填 R9/R10 —— fill-down 缺失。"""
        missing_g = [r for r in _DATA_ROWS if "<f" not in cells.get(f"G{r}", "")]
        missing_i = [r for r in _DATA_ROWS if "<f" not in cells.get(f"I{r}", "")]
        assert missing_g == [10, 11, 12, 13]
        assert missing_i == [11, 12, 13]
        by_where = {r[1]: r for r in TEMPLATE_FORMULA_COVERAGE_DEFECTS_G1202}
        assert "G10:G13" in by_where and "I11:I13" in by_where
        assert "前端 fvCheck 落库补齐" in by_where["G10:G13"][3]
        assert "前端 netHedgePnl 落库补齐" in by_where["I11:I13"][3]

    def test_defect_5_footer_misses_the_h_column_total(self, cells, ws) -> None:
        """🔴 缺陷⑤：`H14` 整格无公式 —— footer 漏了 H 列合计（八列都有）。"""
        assert "<f" not in cells.get(f"H{FOOTER_ROW_G1202}", "")
        assert ws[f"H{FOOTER_ROW_G1202}"].value is None
        assert "H" not in TEMPLATE_FOOTER_FORMULAS_G1202
        assert sorted(TEMPLATE_FOOTER_FORMULAS_G1202) == [
            "B", "C", "D", "E", "F", "G", "I", "J",
        ]
        entry = next(r for r in TEMPLATE_FORMULA_COVERAGE_DEFECTS_G1202 if r[1] == "H14")
        assert entry[3] == "应为 =SUM(H9:H13)"
        # H 列在数据区是有值的（R10=-240000）⇒ 漏合计是真缺陷而不是「本列不该有合计」
        assert ws["H10"].value == -240000

    @pytest.mark.parametrize("col", sorted(TEMPLATE_FOOTER_FORMULAS_G1202))
    def test_footer_formula_ledger_matches_template_cell_by_cell(self, col: str, ws) -> None:
        got = ws[f"{col}{FOOTER_ROW_G1202}"].value
        assert got == TEMPLATE_FOOTER_FORMULAS_G1202[col], (
            f"{col}14 模板实得 {got!r}，台账为 {TEMPLATE_FOOTER_FORMULAS_G1202[col]!r}"
        )

    def test_footer_shared_group_is_interrupted(self, cells) -> None:
        """🔴 `F14` 是 shared 主格、ref=`F14:J14` 覆盖五列，但 `H14` 无公式、`I14` 独立公式
        ⇒ 实际成员只有 `G14`/`J14` 两格。

        按「组 ref 覆盖的列都是成员」去验会有两格假红；按「footer 全列同形态」去验会在
        H14 与 I14 各打一次红。
        """
        facts = FOOTER_SHAPE_FACTS_G1202
        assert facts["shared_master"] == "F14"
        assert re.search(r'<f t="shared" ref="F14:J14" si="(\d+)">', cells["F14"]), (
            f"F14 不再是 ref=F14:J14 的 shared 主格：{cells['F14'][:70]!r}"
        )
        si = re.search(r'si="(\d+)"', cells["F14"]).group(1)  # type: ignore[union-attr]
        for ref in facts["shared_members"]:  # type: ignore[union-attr]
            assert cells[str(ref)].startswith(f'<f t="shared" si="{si}"/>'), (
                f"{ref} 不再是 si={si} 的成员格：{cells[str(ref)][:60]!r}"
            )
        # 打断点：H14 无公式、I14 是独立 plain
        assert "<f" not in cells["H14"]
        assert "<f>" in cells["I14"] and 't="shared"' not in cells["I14"]
        assert facts["no_formula_columns"] == ["A", "H"]
        assert facts["plain_sum_columns"] == ["B", "C", "D", "E", "I"]

    def test_footer_carries_total_formula_is_declared(self, ws) -> None:
        assert SPEC_G1202.footer_carries_total_formula is True
        assert str(ws[f"A{FOOTER_ROW_G1202}"].value) == SPEC_G1202.footer_marker == "合计"


# ════════════════════════════════════════════════════════════════════════════
# ⑤ 几何 / 零裸 IF / store 读写
# ════════════════════════════════════════════════════════════════════════════
class TestGeometryAndStore:
    def test_single_region_r9_to_r13_with_only_two_seeded_rows(self, ws) -> None:
        """🔴 模板只预填 R9/R10，R11-R13 整行全空（样式预留）。

        `last_data_row=13` 的依据是 footer 的 SUM 区间 `x9:x13` —— 模板自己把数据区画到 R13。
        """
        assert (SPEC_G1202.first_data_row, SPEC_G1202.last_data_row) == (9, 13)
        assert len(_DATA_ROWS) == 5
        assert SEEDED_DATA_ROWS_G1202 == (9, 10)
        assert BLANK_DATA_ROWS_G1202 == (11, 12, 13)
        assert set(SEEDED_DATA_ROWS_G1202) | set(BLANK_DATA_ROWS_G1202) == set(_DATA_ROWS)
        for r in BLANK_DATA_ROWS_G1202:
            vals = [ws.cell(row=r, column=c).value for c in range(1, 11)]
            assert all(v in (None, "") for v in vals), f"R{r} 不再是空行：{vals}"
        for r in SEEDED_DATA_ROWS_G1202:
            assert ws.cell(row=r, column=1).value not in (None, ""), f"R{r} 的 A 列空了"
        # footer 的 SUM 区间确实覆盖到 last_data_row
        assert TEMPLATE_FOOTER_FORMULAS_G1202["C"] == "=SUM(C9:C13)"
        assert SPEC_G1202.footer_row == FOOTER_ROW_G1202 == 14
        assert SPEC_G1202.store_item_id == STORE_ITEM_ID_G1202 == "G12-hedge-detail-rows"
        assert not SPEC_G1202.row_section_field

    def test_audit_note_rows_below_footer_are_out_of_scope(self, ws) -> None:
        assert str(ws["A15"].value or "").startswith("三、审计说明")
        assert str(ws["E15"].value or "").startswith("四、审计结论")
        assert 15 > SPEC_G1202.footer_row

    def test_managed_sheet_has_zero_bare_if(self, neutralized) -> None:
        """🔴 受管 sheet 零裸 IF（整册 7 格全在别的 sheet）⇒ 中性化不动本表。"""
        _after, touched = neutralized
        mine = [t for t in touched if SHEET_XML_G1202.split("/")[-1] in t]
        assert mine == [], f"受管 sheet 被中性化命中 {len(mine)} 格"
        assert touched, "整册应仍有裸 IF（别的 sheet）—— 全零说明中性化没跑"

    def test_neutralization_leaves_the_three_existing_formulas_intact(
        self, neutralized
    ) -> None:
        after, _touched = neutralized
        for (col, row) in TEMPLATE_ROW_FORMULAS_G1202:
            assert "<f" in after[f"{col}{row}"], f"{col}{row} 中性化后丢了公式"

    def test_row_identity_key_is_rowid(self) -> None:
        assert SPEC_G1202.row_identity_key == "rowId"
        assert _frontend_row_fields()[0] == "rowId"

    def test_iter_reads_every_row_by_rowid(self) -> None:
        payload = [{"rowId": "g12h-a"}, {"rowId": "g12h-b"}]
        assert [rid for rid, _ in iter_store_rows(SPEC_G1202, payload)] == [
            "g12h-a",
            "g12h-b",
        ]

    def test_duplicate_rowid_fails_closed(self) -> None:
        from app.services.workpaper_sync.phase5_row_table_sheet import (
            RowTableStorePayloadError,
        )

        with pytest.raises(RowTableStorePayloadError, match="重复行身份"):
            list(iter_store_rows(SPEC_G1202, [{"rowId": "x"}, {"rowId": "x"}]))

    def test_merge_keys_rows_by_rowid_including_the_derived_columns(self) -> None:
        """🔴 `G`/`I` 判 `editable` ⇒ 它们**会**进 store（与别家的公式列相反）。"""

        class _FV:
            def __init__(self, key: str, row_key: str, value):
                self.stable_key = key
                self.row_key = row_key
                self.value = value
                self.is_protected = False

        keys = {
            f"{SPEC_G1202.table_key}/g12h-a/sales_portion": 1000,
            f"{SPEC_G1202.table_key}/g12h-a/net_hedge_pnl": 1000,
            f"{SPEC_G1202.table_key}/g12h-a/fv_check": True,
        }

        class _Proj:
            def stable_keys(self):
                return list(keys)

            def get(self, k):
                return _FV(k, "g12h-a", keys[k]) if k in keys else None

        rows, applied, _visited, touched = merge_projection_into_store_rows(
            SPEC_G1202,
            projection=_Proj(),
            base_rows=[{"rowId": "g12h-a"}, {"rowId": "g12h-b"}],
        )
        assert applied == 3 and touched == {"g12h-a"}
        hit = next(r for r in rows if r["rowId"] == "g12h-a")
        assert hit["salesPortion"] == 1000
        assert hit["netHedgePnl"] == 1000
        assert hit["fvCheck"] is True
        assert [r["rowId"] for r in rows] == ["g12h-a", "g12h-b"]

"""真库 D4 entry（39 binding）物化 harness —— 基线测试与证据脚本共用同一份。

spec: oo-single-pass-materialize-and-room-leave · Requirement 1.1 / 1.5

阶段 1 要同时交付两样东西：

① 把「39 趟」钉成判据的**基线测试**（`test_single_pass_materialize.py`）；
② 单次物化的墙钟耗时 + 内存峰值**证据**（`scripts/analyze/measure_d4_materialize_baseline.py`）。

两者必须跑**同一个** world：各自拼一份 binding 列表就是两个不同的 39，测的与记的就不是
同一件事了。因此 world 的构造与调用计数都只有这一份实现。

harness 刻意复用生产 attach/publish 的同一内核（`_align_specs_to_sibling_tables` /
`_static_region_bindings`）而不是手写 binding 表：binding 数（39）是本 spec 的核心量，
它必须来自真实契约。
"""
from __future__ import annotations

import contextlib
import dataclasses
import functools
import hashlib
import os
import sys
import traceback
import uuid
from collections import Counter
from dataclasses import dataclass
from pathlib import Path
from typing import Any, Callable, Iterator, Mapping

_REPO = Path(__file__).resolve().parents[3]
_BACKEND = _REPO / "backend"
if str(_BACKEND) not in sys.path:  # pragma: no cover - import 环境自举
    sys.path.insert(0, str(_BACKEND))
os.environ.setdefault("DB_DISABLE_SSL", "True")

from app.services.excel_structure_fingerprint import (  # noqa: E402
    GT_SYNC_SHEET_NAME,
    identity_inventory,
)
from app.services.workpaper_sync import excel_instrumentation as EI  # noqa: E402
from app.services.workpaper_sync import excel_materialize as EM  # noqa: E402
from app.services.workpaper_sync import phase5_d4_revenue_detail as D4  # noqa: E402
from app.services.workpaper_sync.adapters import excel as AX  # noqa: E402
from app.services.workpaper_sync.adapters.base import ValueType  # noqa: E402
from app.services.workpaper_sync.contracts import parse_contract  # noqa: E402
from app.services.workpaper_sync.excel_entry_gate import (  # noqa: E402
    AdapterBuild,
    FrozenEntryDefinitions,
    parse_identity_inventory,
)
from app.services.workpaper_sync.excel_extract import ExcelIdentityBinding  # noqa: E402
from app.services.workpaper_sync.models import (  # noqa: E402
    AuthorityModel,
    BundleSlot,
    BundleSlotSpec,
    DefinitionState,
)
from app.services.workpaper_sync.projection_first_publication import (  # noqa: E402
    _align_specs_to_sibling_tables,
    _row_bearing_table_key,
    _static_region_bindings,
)
from app.services.workpaper_sync.resolution import (  # noqa: E402
    DefinitionBundleSnapshot,
)

__all__ = [
    "D4World",
    "ENTRY_ID",
    "MaterializeCallCounter",
    "build_world",
    "force_chained_path",
    "rebased_world",
]

ENTRY_ID = D4.ENTRY_ID


def _digest(label: str) -> str:
    return hashlib.sha256(label.encode("utf-8")).hexdigest()


def _bundle(contract: Any) -> DefinitionBundleSnapshot:
    def slot(kind: BundleSlot, digest: str) -> BundleSlotSpec:
        return BundleSlotSpec(kind, "definition", f"definition:{uuid.uuid4()}", digest)

    return DefinitionBundleSnapshot(
        bundle_id=uuid.uuid4(),
        bundle_sha256=_digest("single-pass-bundle"),
        schema_version="definition-bundle:v1",
        state=DefinitionState.approved,
        authority_model=AuthorityModel.projection_contract,
        authority_model_definition_id=uuid.uuid4(),
        authority_model_definition_sha256=_digest("single-pass-authority"),
        slots={
            BundleSlot.template: slot(
                BundleSlot.template, contract.template_definition_sha256
            ),
            BundleSlot.instrumentation: slot(
                BundleSlot.instrumentation, contract.instrumentation_definition_sha256
            ),
            BundleSlot.contract: slot(BundleSlot.contract, contract.canonical_sha256),
        },
    )


@dataclass(frozen=True)
class _FrozenD4:
    """真实契约 / instrumented 字节 / 冻结定义 / binding 集合（构造一次、复用）。"""

    contract: Any
    instrumented_bytes: bytes
    definitions: FrozenEntryDefinitions
    primary: ExcelIdentityBinding
    siblings: tuple[ExcelIdentityBinding, ...]

    @property
    def binding_count(self) -> int:
        return 1 + len(self.siblings)


@functools.lru_cache(maxsize=1)
def _frozen_d4() -> _FrozenD4:
    """真 D4 模板 + 真 instrumentation + 真契约（不是合成 fixture）。

    instrumentation 要跑一遍真模板（实测 199KB 压缩 / 1.3MB 解压 / 46 sheet —— 既有注释
    里「~7MB 工作簿」的说法**与实测不符**，39 趟的代价来自重复解析而不是文件大），
    耗时在秒级 ⇒ 进程内只做一次。
    """
    D4.assert_contract_file_matches_source()
    contract = parse_contract(D4.build_contract_payload(), adapter_id=D4.ADAPTER_ID)
    instrumented = EI.instrument_workbook_bytes_multi(
        D4.read_authoritative_template(),
        D4.instrumentation_specs(),
        gate=D4.excel_carrier_gate(),
    )
    inventory = identity_inventory(
        instrumented.instrumented_bytes,
        expected_table=D4.TABLE_NAME,
        uuid_column_letter=D4.UUID_COL,
    )
    definitions = FrozenEntryDefinitions(
        entry_id=D4.ENTRY_ID,
        bundle=_bundle(contract),
        contract=contract,
        adapter_build=AdapterBuild(
            adapter_id=D4.ADAPTER_ID,
            adapter_build_digest=_digest("single-pass-adapter"),
            document_type="xlsx",
            contract_version=contract.semantic_version,
        ),
        identity_inventory=parse_identity_inventory(inventory),
        business_sheets=(),
        dynamic_column_keys={},
        structure_inventory_size=0,
    )
    primary, siblings = _bindings(contract)
    return _FrozenD4(
        contract=contract,
        instrumented_bytes=instrumented.instrumented_bytes,
        definitions=definitions,
        primary=primary,
        siblings=siblings,
    )


def _bindings(
    contract: Any,
) -> tuple[ExcelIdentityBinding, tuple[ExcelIdentityBinding, ...]]:
    """primary + sibling binding，规则与生产 attach/publish 路径**同一内核**。

    刻意复用 `_align_specs_to_sibling_tables` / `_static_region_bindings`：自己拼一份
    binding 列表等于在 harness 里造一个与生产不同的 39，那样「趟数」这条判据就与真实
    路径脱钩了。
    """
    primary = ExcelIdentityBinding(
        table_name=D4.TABLE_NAME,
        uuid_column=D4.UUID_COL,
        table_key=_row_bearing_table_key(contract=contract, provider=D4),
        metadata_sheet=GT_SYNC_SHEET_NAME,
        defined_name_prefix="GT_",
        tombstoned_row_keys=(),
        dynamic_column_columns={},
    )
    siblings = [
        ExcelIdentityBinding(
            table_name=str(spec.table_name),
            uuid_column=str(spec.uuid_col),
            table_key=str(dynamic.table_key),
            metadata_sheet=GT_SYNC_SHEET_NAME,
            defined_name_prefix=str(getattr(spec, "defined_name_prefix", None) or "GT_"),
            tombstoned_row_keys=(),
            dynamic_column_columns={},
        )
        for spec, dynamic in _align_specs_to_sibling_tables(
            provider=D4, contract=contract, primary=primary
        )
    ]
    siblings.extend(
        _static_region_bindings(provider=D4, metadata_sheet=GT_SYNC_SHEET_NAME)
    )
    return primary, tuple(siblings)


def _text_safe(projection: Any) -> Any:
    """把 `text` 字段里反读出来的非字符串值转成字符串。

    模板里有几个受管 text 格存的是数字（`d4_32_groups/*/group_label` 的 `1`），extract
    如实返回 int，而 materialize 的 `text` 规范化拒绝 int（「不得把数字与其字符串形态
    判等」）。那是**模板内容**与契约类型的既有不一致，与本 spec 无关；harness 在入口把它
    规整掉，免得一条与写入趟数无关的 `EditableCellWriteError` 挡住基线判据。
    """
    values = {
        key: (
            dataclasses.replace(field, value=str(field.value))
            if field.value_type is ValueType.text
            and field.value is not None
            and not isinstance(field.value, str)
            else field
        )
        for key, field in projection.values.items()
    }
    return dataclasses.replace(projection, values=values)


@dataclass(frozen=True)
class D4World:
    """一次物化所需的全部真实输入。"""

    adapter: Any
    base: Path
    staged_dir: Path
    projection: Any
    contract: Any
    binding_count: int

    def staged(self, name: str) -> Path:
        return self.staged_dir / name

    def materialize(self, name: str = "staged.xlsx") -> Path:
        """跑一次完整 `adapter.materialize`，返回产物路径。"""
        output = self.staged(name)
        self.adapter.materialize(
            substrate=self.base,
            projection=self.projection,
            output=output,
            contract=self.contract,
        )
        return output


def build_world(workdir: Path) -> D4World:
    """在 `workdir` 下铺开 substrate 与 staging 命名空间，返回可直接物化的 world。

    projection 取自 substrate **自身**的反读结果：它必然覆盖**每个** binding 的受管区，
    是「全部 binding 都被写一遍」的最诚实输入（合成 projection 只会命中其中几张表，
    那样趟数就不是 39 了）。
    """
    frozen = _frozen_d4()
    workdir.mkdir(parents=True, exist_ok=True)
    base = workdir / "d4-base.xlsx"
    base.write_bytes(frozen.instrumented_bytes)
    staged_dir = workdir / AX.STAGING_NAMESPACE
    staged_dir.mkdir(parents=True, exist_ok=True)
    adapter = AX.build_excel_adapter(
        definitions=frozen.definitions,
        binding=frozen.primary,
        sibling_bindings=frozen.siblings,
        direction="html_to_oo",
    )
    projection = _text_safe(
        adapter.extract(artifact=base, contract=frozen.contract)
    )
    return D4World(
        adapter=adapter,
        base=base,
        staged_dir=staged_dir,
        projection=projection,
        contract=frozen.contract,
        binding_count=frozen.binding_count,
    )


def rebased_world(world: D4World, *, substrate_bytes: bytes, label: str) -> D4World:
    """以给定字节为 substrate 重挂一个 world（projection 取新 substrate 的**自反读**）。

    主用途是 **steady-state substrate**：生产的 substrate 恒是「上一次发布的产物」（已被
    openpyxl 归一化过），而 `build_world` 铺的是 instrumented **模板** —— 两者在 verify
    眼里不是一回事。任务 7（design 附录 E.2）与任务 8（附录 F.8 第 2 条）都实测过这个差别：
    模板 substrate 上 `verify_unmanaged_regions` 39/39 判漂移、首个差异恒为
    `workbook_and_styles`（模板的 `styles.xml` 从未被 openpyxl 重序列化，而转置 sheet 的
    `wb.save()` 会重写整簿状态）。要在 verify 的**结论**上做变异反证，分母必须是 steady
    state，否则「必红」在第一代 world 上本来就红，反证没有信号。

    典型用法（第二代 = steady state）::

        gen1 = world.materialize("gen1.xlsx")
        steady = rebased_world(world, substrate_bytes=gen1.read_bytes(), label="steady")

    `adapter` / `contract` / `binding_count` 原样沿用（substrate 路径是**每次调用**的入参，
    不挂在 adapter 上），只换 `base` / `staged_dir` / `projection` 三项。
    """
    root = world.base.parent / label
    root.mkdir(parents=True, exist_ok=True)
    base = root / "d4-base.xlsx"
    base.write_bytes(substrate_bytes)
    staged_dir = root / AX.STAGING_NAMESPACE
    staged_dir.mkdir(parents=True, exist_ok=True)
    projection = _text_safe(
        world.adapter.extract(artifact=base, contract=world.contract)
    )
    return dataclasses.replace(
        world, base=base, staged_dir=staged_dir, projection=projection
    )


class MaterializeCallCounter:
    """统计一次物化里的**真实** openpyxl 读写次数与链式趟数。

    三条计数各有各的理由：

    * ``loads``（`openpyxl.load_workbook`）—— 需求 1.1 的左半；打在 `load_workbook` 上而
      不是 `_acquire_read_only_workbook` 上，因为后者命中作用域缓存时并不解析，统计它会把
      「共享了一次解析」也算成一次，于是 scope 复用的效果在判据里看不见。
    * ``saves``（`Workbook.save`）—— 需求 1.1 的右半。打在**类**上而不是某个模块属性上，
      这样不管调用方怎么拿到 workbook 都会被计到。
    * ``trips``（`materialize_projection` **与** `materialize_projection_single_pass`）——
      就是「39 趟」这个量本身。两个入口都计进同一个 ``trips``，因为「趟」问的是「这次物化
      做了几次写入 pass」：链式 39 次、单趟 1 次，同一把尺子量两条路径才比得出来。
      单趟入口刻意**只在返回后**才计数：decline 会抛 `SinglePassDeclined`，它一个字节都没写，
      把它算成一趟会让「趟数」在回落时变成 40，那是假的。

    ``load_workbook`` 的打桩刻意扫 `sys.modules` 重绑所有 `from openpyxl import
    load_workbook` 的模块（`phase5_transposed_sheet` 就是这么导的）：只改
    `openpyxl.load_workbook` 会漏掉它们，把基线记少 —— 基线记少了，后面的提速就会被
    高估。
    """

    def __init__(self) -> None:
        self.loads: list[str] = []
        self.saves: list[str] = []
        self.trips: list[str] = []
        #: 其中走**单趟**入口的那些趟（成功返回才入册）。拆出来是为了让判据能说清
        #: 「趟数 1」是单趟命中，而不是「链式恰好只有一个 binding」。
        self.single_pass_trips: list[str] = []
        #: 调用点 → 次数。「83」这个总数本身说明不了问题，拆开才看得出哪一项是趟数带来的：
        #: 78 在 `_acquire_read_only_workbook`（39 趟 × data_only 两视图），另 5 次是转置
        #: sheet 与整簿指纹的固定开销 —— 后者不随 binding 数变，单趟化也带不走。
        self.load_sites: Counter[str] = Counter()
        self.save_sites: Counter[str] = Counter()
        #: 落在**文件路径**上的解析 = substrate 链（base + 链式中间产物）。需求 1.1 说的
        #: 「一次物化对 substrate 的 load_workbook 次数」就是这一项；BytesIO 上的那几次
        #: （转置 sheet / 整簿指纹）是与 binding 数无关的固定开销，混在总数里会把判据糊掉。
        self.substrate_loads: list[str] = []
        self.in_memory_loads: list[str] = []

    @property
    def load_count(self) -> int:
        return len(self.loads)

    @property
    def save_count(self) -> int:
        return len(self.saves)

    @property
    def trip_count(self) -> int:
        return len(self.trips)

    @property
    def single_pass_trip_count(self) -> int:
        return len(self.single_pass_trips)

    @property
    def substrate_load_count(self) -> int:
        return len(self.substrate_loads)

    @property
    def in_memory_load_count(self) -> int:
        return len(self.in_memory_loads)

    def as_dict(self) -> Mapping[str, Any]:
        return {
            "load_workbook": self.load_count,
            "load_workbook_on_files": self.substrate_load_count,
            "load_workbook_in_memory": self.in_memory_load_count,
            "workbook_save": self.save_count,
            "materialize_trips": self.trip_count,
            "single_pass_trips": self.single_pass_trip_count,
            "load_workbook_by_site": dict(self.load_sites.most_common()),
            "workbook_save_by_site": dict(self.save_sites.most_common()),
        }

    @staticmethod
    def _caller_site() -> str:
        """最内层**非 harness** 的调用帧，形如 `excel_extract.py:3087 <函数名>`。"""
        for frame in reversed(traceback.extract_stack()):
            if Path(frame.filename).name != Path(__file__).name:
                return f"{Path(frame.filename).name}:{frame.lineno} {frame.name}"
        return "<unknown>"

    @contextlib.contextmanager
    def installed(self) -> Iterator["MaterializeCallCounter"]:
        """安装计数桩；退出时**逐一**还原（脚本与测试共用，不依赖 monkeypatch）。"""
        import openpyxl
        from openpyxl.workbook.workbook import Workbook

        undo: list[Callable[[], None]] = []
        original_load = openpyxl.load_workbook

        def counted_load(*args: Any, **kwargs: Any) -> Any:
            target = args[0] if args else kwargs.get("filename")
            label = str(target)
            self.loads.append(label)
            self.load_sites[self._caller_site()] += 1
            if isinstance(target, (str, os.PathLike)):
                self.substrate_loads.append(label)
            else:
                self.in_memory_loads.append(label)
            return original_load(*args, **kwargs)

        # openpyxl 本体 + 所有 `from openpyxl import load_workbook` 的模块。
        targets = [openpyxl] + [
            module
            for module in list(sys.modules.values())
            if module is not None
            and getattr(module, "load_workbook", None) is original_load
        ]
        for module in targets:
            undo.append(
                lambda m=module, f=original_load: setattr(m, "load_workbook", f)
            )
            setattr(module, "load_workbook", counted_load)

        original_save = Workbook.save

        def counted_save(wb: Any, filename: Any) -> Any:
            self.saves.append(str(filename))
            self.save_sites[self._caller_site()] += 1
            return original_save(wb, filename)

        undo.append(lambda: setattr(Workbook, "save", original_save))
        Workbook.save = counted_save  # type: ignore[method-assign]

        original_trip = AX.materialize_projection

        def counted_trip(*args: Any, **kwargs: Any) -> Any:
            binding = kwargs.get("binding")
            self.trips.append(str(getattr(binding, "table_key", binding)))
            return original_trip(*args, **kwargs)

        undo.append(lambda: setattr(AX, "materialize_projection", original_trip))
        AX.materialize_projection = counted_trip  # type: ignore[assignment]

        # 单趟入口。打在 `excel_materialize` 模块属性上而不是 adapter 上：
        # `_try_single_pass_materialize` 在**函数体内**做 `from … import
        # materialize_projection_single_pass`，每次调用都重新取模块属性 ⇒ 打模块就打得到。
        original_single = EM.materialize_projection_single_pass

        def counted_single(*args: Any, **kwargs: Any) -> Any:
            outcome = original_single(*args, **kwargs)
            # 成功返回才算一趟：decline 抛异常、一个字节都没写，不该计数。
            label = f"single_pass[{len(kwargs.get('bindings') or ())} binding]"
            self.trips.append(label)
            self.single_pass_trips.append(label)
            return outcome

        undo.append(
            lambda: setattr(EM, "materialize_projection_single_pass", original_single)
        )
        EM.materialize_projection_single_pass = counted_single  # type: ignore[assignment]

        try:
            yield self
        finally:
            for restore in reversed(undo):
                restore()


@contextlib.contextmanager
def force_chained_path() -> Iterator[None]:
    """强制走**逐 binding 链式**写盘路径（= 单趟 decline 时的**回落**路径）。

    任务 3 落地后单趟对真 D4 已生效（默认路径趟数 1），但链式路径**不是**死代码：它是
    decline 三条（结构性插行 / openpyxl 全量重写 / 同格 payload 冲突）的正式回落路径，由
    `_try_single_pass_materialize` 返回 `None` 时真实到达。真库上任一 entry 一旦插行就会走它。

    因此本 helper 继续存在，作用是把回落路径**显式**量出来：判据不能建立在「单趟恰好
    decline」这个偶然上（那是任务 3 之前的状态），要钉的是「链式实现 = 一 binding 一趟」
    这条结构事实。用它的那条判据见 `test_chained_fallback_path_call_counts`。
    """
    if not hasattr(AX.ExcelSyncAdapter, "_try_single_pass_materialize"):
        # 链式已是唯一路径（任务 3 之前的 HEAD 就是这样）⇒ 无需摘任何东西。
        yield
        return
    original = AX.ExcelSyncAdapter._try_single_pass_materialize
    AX.ExcelSyncAdapter._try_single_pass_materialize = (  # type: ignore[method-assign]
        lambda self, **kwargs: None
    )
    try:
        yield
    finally:
        AX.ExcelSyncAdapter._try_single_pass_materialize = original  # type: ignore[method-assign]

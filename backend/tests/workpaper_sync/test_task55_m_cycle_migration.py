# -*- coding: utf-8 -*-
r"""Task 55 守卫 —— M 循环（权益循环）Excel 独立 entry 逐一迁移。

spec: workpaper-html-onlyoffice-bidirectional-writeback-closure · Wave 5 Task 55
Properties: 24 / 28 / 69 / 70
Requirements: 6.6 · 6.10 · 12.1 · 12.4 · 12.10 · 12.11 · 12.12 · 14.1

═══ 本文件守什么 ═══

两份产物（`workpaper_sync_m_cycle_manifest_slice.json` /
`workpaper_sync_m_cycle_deletion_plan.json`）里的**每个数字与每条结论**都必须能从磁盘现算
出来。判据一律落在「行为 / 结构 / 真实执行」上，不是「字符串存在」：

* 消费边 = statement-position import **路径**字面量，且排除双引号字符串内的匹配
  （`workpaperSyncLegacyBaseline.generated.ts` 的 JSON `"snippet"` 字段含完整 `from '...'`
  形态；不排除会把死桩算成活的）。
* orphan 判定用**可达性**而非入度；barrel 入边用路径解析而非 stem 相等。
* 剥注释**保留行号**（同长空白替换）。删行式剥注释会让全部 `#Lnn` 锚点整体上移 ⇒ 假红。
* 「某声明是否真生效」落到**模板形态判据**（遍历 + 外层门控 + 内层嵌套三要素）。
* 声明与现算**两侧都验**：只断言「现算为 0」时，「声明写 1、实际 0」这种谎话能通过
  （Task 54 的 M20 教训）。

═══ M 循环相对前九轮的八处形态差异（照抄 L / K 的判据必假红）═══

1. **dual-mode 载体成对孪生**：`useM{n}DualMode.ts`（自带完整实现）+
   `useM{n}EntryDualMode.ts`（薄封装共享基类）各 10 个 = 20 个文件 == entry 数 × 2。
   orphan 11 / live 9 —— **M9 的两个孪生都是 orphan**。L 是 9 文件 / 8 entry 三分，K 是 13 == 13 二分。
2. 🔴 **inert 开关 0 条**：9 条可兑现 + 1 条完全无开关（M9）。L 有 4 条 inert，
   抄 L 那份 `switch_present_but_inert` 会直接判错。本文件用「三要素齐验」+ 合成 inert 样本
   证明扫描器不是恒判 redeemable。
3. 🔴 **活载体的 `M{n}_SHEET_MAP` 有 11 对指向权威册不存在的 sheet 名**（84 对里）。
   前九轮从不核这一层 —— 那是 AC 6.10「结构漂移 fail closed」在 sheet-name 层的缺口。
   反向判据：10 条 fallback 字面量必须**全部命中**真 sheet（证明比对器不是恒判不存在）。
4. 🔴 **位置化行身份的主形态在持久化键上**：`itemId: \`M{n}-…-row-${n}-…\`` 现算 44 处
   （10 条 entry 全中），而 L/K 那套只看 `rowKey:`/`id:` 值的扫描器只命中 1 处。
   沿用旧扫描器会把 44 处系统性缺陷判成 1 处。
5. **模板与 entry 1:1 且没有 M0**：10 本册 / 10 条 entry，`belongs_to_entry` 无 null。
   L 是 9 本册 / 8 entry（多一本 L0）—— 抄 L 的「差 1」必假红。
6. **活路径无 localStorage 键**：共享基类把 mode 存内存 `ref`，每次挂载重置为 `'html'`。
   键只在 11 个 orphan 里（`m{n}-dual-mode`）。L1 活路径有 `l1-proc:` 键。
7. 🔴 **共享基类的边会收缩**：29 → 19（M 域贡献 10 条）。L 贡献 0 条、清册写「前后不变」。
8. **3 张「修订前」Q 表有三种不同待遇**：M6 专用 `skip-q6a` 分支（正确）、
   M8 静默折叠到 `procedure`、M10 折叠且它那条 `修订前` 判断是**不可达死代码**（排序在后）。

═══ 运行 ═══

    .\.venv\Scripts\python.exe -m pytest backend/tests/workpaper_sync/test_task55_m_cycle_migration.py -q

变异检验脚本：`backend/scripts/diagnose/mutate_task55_m_cycle_guards.py`
"""
from __future__ import annotations

import hashlib
import importlib.util
import itertools
import json
import pathlib
import re
from types import ModuleType
from typing import Any

import pytest
from openpyxl import load_workbook

# ════════════════════════════════════════════════════════════════════════════
# 路径常量
# ════════════════════════════════════════════════════════════════════════════
_THIS = pathlib.Path(__file__).resolve()
ROOT = _THIS.parents[3]
BACKEND = ROOT / "backend"
FRONTEND = ROOT / "audit-platform" / "frontend" / "src"
WP_COMPONENTS = FRONTEND / "components" / "workpaper"
WP_COMPOSABLES = WP_COMPONENTS / "composables"
SRC_COMPOSABLES = FRONTEND / "composables"
CYCLE_COMPOSABLES = SRC_COMPOSABLES / "workpaper"
SYNC_DIR = WP_COMPONENTS / "sync"
DATA = BACKEND / "data"
SVC = BACKEND / "app" / "services" / "workpaper_sync"

MANIFEST_SLICE_PATH = DATA / "workpaper_sync_m_cycle_manifest_slice.json"
DELETION_PLAN_PATH = DATA / "workpaper_sync_m_cycle_deletion_plan.json"
PARADIGM_PATH = DATA / "workpaper_sync_migration_paradigm.json"
FULL_MANIFEST_PATH = DATA / "workpaper_sync_entry_manifest.json"
OVERLAY_PATH = DATA / "workpaper_sync_entry_overlay.json"
SIBLING_SLICE_PATHS = {
    letter: DATA / f"workpaper_sync_{letter.lower()}_cycle_manifest_slice.json"
    for letter in ("D", "E", "F", "G", "H", "I", "J", "K", "L")
}
CONTRACT_DIR = DATA / "workpaper_sync_contracts"
TEMPLATE_DIR = BACKEND / "wp_templates"
TEMPLATE_INDEX = TEMPLATE_DIR / "_index.json"
M_TEMPLATE_DIR = TEMPLATE_DIR / "M"
REFERENCE_COPY_DIR = ROOT / "基础数据" / "致同通用审计程序及底稿模板（2025年修订）"
CHECKLIST_ROUTER = BACKEND / "app" / "routers" / "checklist_responses.py"
REGISTRY = SVC / "adapters" / "registry.py"
HTML_RENDERER_REGISTRY = WP_COMPONENTS / "htmlRendererRegistry.ts"
SHARED_BASE = WP_COMPOSABLES / "useWorkpaperEntryDualMode.ts"
COMPOSABLES_BARREL = WP_COMPOSABLES / "index.ts"
NOTICE_MODULE = SYNC_DIR / "workpaperEntrySyncNotice.ts"
NOTICE_COMPONENT = SYNC_DIR / "GtEntrySyncCapabilityNotice.vue"
NOTICE_COMPONENT_NAME = "GtEntrySyncCapabilityNotice"
CONTRACT_GUARD = BACKEND / "tests" / "workpaper_sync" / "test_migration_paradigm_contract.py"
COVERAGE_GUARD = BACKEND / "tests" / "workpaper_sync" / "test_slice_schema_validator_coverage.py"

CONFLICTS_PY = SVC / "conflicts.py"
MERGE_PY = SVC / "merge.py"
REMATERIALIZE_PY = SVC / "excel_rematerialize.py"
EXTRACT_PY = SVC / "excel_extract.py"
CONTRACTS_PY = SVC / "contracts.py"

CAPABILITY_ENUM = ("bidirectional", "single_html", "single_onlyoffice", "unreachable")
HTML_COUNTERPART_VERDICTS = ("none", "exists")
SWITCH_VERDICTS = ("switch_redeemable", "switch_present_but_inert", "no_switch_at_all")

#: 四条 pilot 契约的 `review.entry_id` 实测值 —— 🔴 按 entry_id 判归属，不按文件名猜。
PILOT_CONTRACT_OWNERS = {
    "xlsx/b60/gt-b60-bundle",
    "xlsx/gt-d2-accounts-receivable",
    "xlsx/gt-g7-long-term-equity-main",
    "xlsx/gt-h1-fixed-assets",
}
#: 第 5 个契约文件是 candidate：`review_status == "candidate"` 且 `review.entry_id` 为 **null**。
#: 它是「candidate 不得注册生产 adapter」的反例分母 —— **不得**要求其 entry_id 非空。
CANDIDATE_CONTRACT_FILE = "_example.candidate.json"

M_CODES = ("M1", "M2", "M3", "M4", "M5", "M6", "M7", "M8", "M9", "M10")
#: 开关**可兑现**的 9 条（三要素齐备）。
REDEEMABLE_CODES = ("M1", "M2", "M3", "M4", "M5", "M6", "M7", "M8", "M10")
#: 完全没有模式切换栏的一条 —— 它是「inert 计数为 0」的对照组。
NO_SWITCH_CODES = ("M9",)
#: dispatch code 折叠的三条。
COLLAPSE_CODES = ("M2", "M8", "M10")


# ════════════════════════════════════════════════════════════════════════════
# 工具
# ════════════════════════════════════════════════════════════════════════════
def _load(path: pathlib.Path) -> dict:
    return json.loads(path.read_text(encoding="utf-8"))


def _sha256_of(path: pathlib.Path) -> str:
    return hashlib.sha256(path.read_bytes()).hexdigest()


def _blank_keep_newlines(match: "re.Match[str]") -> str:
    """把匹配段替换成同长空白，**保留换行** —— 行号绝不能变。"""
    return re.sub(r"[^\n]", " ", match.group(0))


def _strip_comments(source: str) -> str:
    """剥块注释 / 行注释 / HTML 注释，行号保持不变。"""
    source = re.sub(r"/\*.*?\*/", _blank_keep_newlines, source, flags=re.S)
    source = re.sub(r"<!--.*?-->", _blank_keep_newlines, source, flags=re.S)
    source = re.sub(
        r"(?<![:\w\"'`\\])//[^\n]*", lambda m: " " * len(m.group(0)), source
    )
    return source


def _vue_template(source: str) -> str:
    """取 `<template>` 顶层块（不被内层 `<template v-if>` 提前截断），行号保留。"""
    start = source.find("<template>")
    if start < 0:
        return ""
    end = source.rfind("</template>")
    if end < 0:
        return ""
    head = re.sub(r"[^\n]", " ", source[:start])
    return head + source[start:end + len("</template>")]


def _lines_of(text: str, needle: str) -> list[int]:
    return [text.count("\n", 0, m.start()) + 1 for m in re.finditer(re.escape(needle), text)]


def _sheet_names(path: pathlib.Path) -> list[str]:
    """openpyxl 真读 sheet 名。🔴 **不得 strip** —— 源侧多张 sheet 名带前导/尾随空格。"""
    wb = load_workbook(path, read_only=True, data_only=True)
    try:
        return list(wb.sheetnames)
    finally:
        wb.close()


def _formula_cell_counts(path: pathlib.Path, *, data_only: bool = False) -> dict[str, int]:
    """逐 sheet 数公式格。

    🔴 `data_only=False` 是**判据本身**：True 只拿缓存值、公式全丢 ⇒ 计数塌成 0。
    本函数把开关暴露出来，供 `TestGuardSelfChecks` 做反向自检。
    """
    wb = load_workbook(path, data_only=data_only)
    try:
        out: dict[str, int] = {}
        for ws in wb.worksheets:
            n = 0
            for row in ws.iter_rows():
                for cell in row:
                    value = cell.value
                    if isinstance(value, str) and len(value) > 1 and value.startswith("="):
                        n += 1
            if n:
                out[ws.title] = n
        return out
    finally:
        wb.close()


def _indexed_relpaths() -> set[str]:
    idx = _load(TEMPLATE_INDEX)
    return {str(f["relative_path"]).replace("\\", "/") for f in idx["files"]}


def _resolve_repo(rel: str) -> pathlib.Path:
    return ROOT / rel.split("#L")[0]


def _line_no_of(ref: str) -> int:
    m = re.search(r"#L(\d+)$", ref)
    assert m, f"{ref!r} 不含 #L<行号>"
    return int(m.group(1))


def _line_at(ref: str) -> str:
    path = _resolve_repo(ref)
    lines = path.read_text(encoding="utf-8").split("\n")
    n = _line_no_of(ref)
    assert 1 <= n <= len(lines), f"{ref} 行号越界（文件共 {len(lines)} 行）"
    return lines[n - 1]


_IMPORT_FORMS = (
    re.compile(r"""from\s*['"]([^'"\n]+)['"]"""),
    re.compile(r"""import\s*\(\s*['"]([^'"\n]+)['"]"""),
    re.compile(r"""vi\.mock\s*\(\s*['"]([^'"\n]+)['"]"""),
    re.compile(r"""require\s*\(\s*['"]([^'"\n]+)['"]"""),
)


def _inside_double_quoted_string(line: str, pos: int) -> bool:
    """`pos` 是否落在**双引号字符串**内。

    🔴 `workpaperSyncLegacyBaseline.generated.ts` 的 JSON `"snippet"` 字段里含完整的
    ``from '...'`` 形态；不排除它，共享基类的边数结论会被污染。
    """
    depth = 0
    i = 0
    while i < pos and i < len(line):
        ch = line[i]
        if ch == "\\":
            i += 2
            continue
        if ch == '"':
            depth ^= 1
        i += 1
    return bool(depth)


_FE_CACHE: dict[pathlib.Path, str] = {}


def _cached_text(path: pathlib.Path) -> str:
    if path not in _FE_CACHE:
        _FE_CACHE[path] = path.read_text(encoding="utf-8", errors="replace")
    return _FE_CACHE[path]


_FE_FILES: list[pathlib.Path] = []


def _all_frontend() -> list[pathlib.Path]:
    global _FE_FILES
    if not _FE_FILES:
        _FE_FILES = [
            p for p in FRONTEND.rglob("*")
            if p.is_file() and p.suffix in (".ts", ".vue", ".tsx", ".js")
        ]
    return _FE_FILES


def _resolve_spec(spec: str, importer: pathlib.Path) -> pathlib.Path | None:
    """把 import spec 解析成仓库内路径（**路径式**，不是 stem 式）。"""
    if spec.startswith("@/"):
        return FRONTEND / spec[2:]
    if spec.startswith("."):
        return (importer.parent / spec).resolve()
    return None


def _is_test_path(path: pathlib.Path) -> bool:
    p = path.as_posix()
    return "__tests__" in p or p.endswith((".spec.ts", ".test.ts"))


def _statement_edges_to(target: pathlib.Path) -> tuple[list[str], list[str]]:
    """返回 (生产边, 测试边)，只认 statement-position 且排除双引号字符串内的匹配。"""
    prod: list[str] = []
    test: list[str] = []
    stem = target.with_suffix("")
    for f in _all_frontend():
        text = _cached_text(f)
        if target.stem not in text:
            continue
        for i, line in enumerate(text.split("\n"), 1):
            for rx in _IMPORT_FORMS:
                for m in rx.finditer(line):
                    if _inside_double_quoted_string(line, m.start()):
                        continue
                    r = _resolve_spec(m.group(1), f)
                    if r is None:
                        continue
                    if r in (stem, target) or r.with_suffix("") == stem:
                        ref = f"{f.relative_to(ROOT).as_posix()}#L{i}"
                        (test if _is_test_path(f) else prod).append(ref)
    return sorted(prod), sorted(test)


def _wide_scope_edge_files(token: str) -> set[str]:
    """宽口径：逐文件 token 命中即算（用于证明窄口径不是空谈）。"""
    return {
        f.relative_to(ROOT).as_posix()
        for f in _all_frontend() if token in _cached_text(f)
    }


def _dual_mode_module_files() -> list[pathlib.Path]:
    return sorted(FRONTEND.rglob("useM*DualMode.ts"))


def _m_cycle_files() -> list[pathlib.Path]:
    """M 域生产文件（口径与 slice 的扫描分母逐字一致）。"""
    out: list[pathlib.Path] = []
    for n in range(1, 11):
        d = WP_COMPONENTS / f"m{n}"
        if d.is_dir():
            for p in d.rglob("*"):
                if p.is_file() and p.suffix in (".ts", ".vue") and "__tests__" not in p.as_posix():
                    out.append(p)
    for p in sorted(WP_COMPONENTS.iterdir()):
        if p.suffix == ".vue" and re.match(r"^GtM\d+", p.name):
            out.append(p)
    for p in sorted(WP_COMPOSABLES.iterdir()):
        if p.suffix == ".ts" and re.match(r"^(use)?[mM](10|[1-9])(?![0-9])", p.name):
            out.append(p)
    for p in sorted(SRC_COMPOSABLES.iterdir()):
        if p.is_file() and p.suffix == ".ts" and re.match(r"^(use)?[mM](10|[1-9])(?![0-9])", p.name):
            out.append(p)
    for n in range(1, 11):
        d = CYCLE_COMPOSABLES / f"m{n}"
        if d.is_dir():
            for p in d.rglob("*.ts"):
                if "__tests__" not in p.as_posix():
                    out.append(p)
    return sorted(set(out))


# ── 位置化行身份：两族分开 ────────────────────────────────────────────────
_IDENTITY_KEY = re.compile(r"(?<![\w$])(rowId|rowKey|id)\s*:\s*([^,\n]+)")
_POSITIONAL_TOKEN = re.compile(r"(?:^|[^\w$])(?:i|idx|index)(?:\s*\+\s*1)?(?:\s*\}|\s*[,)\]`]|$)")
_POSITIONAL_INTERP = re.compile(r"\$\{\s*(?:i|idx|index)\s*\}")
#: 🔴 M 循环的**主形态**：位置化发生在持久化键模板里，不在行对象的身份字段上。
_PERSIST_POSITIONAL = re.compile(r"itemId:\s*`([^`]*row-\$\{\s*(?:n|i|idx|index)\s*\}[^`]*)`")
_N_IS_INDEX = re.compile(r"const\s+n\s*=\s*(?:i|idx|index)\s*\+\s*1")
_DISPLAY_SEQ_SITE = re.compile(r"\bseq\s*:\s*[^,]*(?:idx|index|\bi\b)\s*\+\s*1")
_ENTROPY_ROW_KEY = re.compile(r"const\s+key\s*=\s*`([a-z0-9-]+)-\$\{Date\.now\(\)\}")
_CALC_ASSIGN = re.compile(r"row\.(\w+)\s*=\s*(?:calc\w+|compute\w+)\(")
_PERSIST_FIELD = re.compile(r"itemId:\s*`[^`]*`,\s*data:\s*\{\s*remark:\s*String\(row\.(\w+)\)")
_SUMMARY_ITEM = re.compile(r"itemId:\s*'([^']*(?:summary|subtotal|total)[^']*)'")

_HARDCODED_PATTERNS = {
    "blankRows_literal_count": re.compile(r"blankRows\s*\([^,)]*,\s*\d+\s*\)"),
    "horizontal_company_column_literals": re.compile(r"['\"]公司\s*[1-9]\d*['\"]"),
    "column_key_uses_label": re.compile(r"key\s*:\s*(?:row|col|c)\.label"),
    "row_cell_key_uses_label": re.compile(r"cells\s*\[\s*(?:row|r)\.label\s*\]"),
    "seed_placeholder_literal": re.compile(r"['\"](?:项目|成本项目|单位)\s*N['\"]"),
    "positional_row_id_template": re.compile(
        r"`(?:seed|row|detail|item)-\$\{\s*(?:i|idx|index)\s*\}"
    ),
}


def _render_key_positional_hits(files: list[pathlib.Path]) -> list[dict[str, str]]:
    """渲染键族（L / K 那套口径）。"""
    hits: list[dict[str, str]] = []
    for p in files:
        src = _strip_comments(_cached_text(p))
        rel = p.relative_to(ROOT).as_posix()
        for i, line in enumerate(src.split("\n"), 1):
            for m in _IDENTITY_KEY.finditer(line):
                val = m.group(2).strip()
                if _POSITIONAL_TOKEN.search(val) or _POSITIONAL_INTERP.search(val):
                    hits.append({"site": f"{rel}#L{i}", "key": m.group(1), "value": val})
    return hits


def _persist_key_positional_hits(files: list[pathlib.Path]) -> list[dict[str, str]]:
    """持久化键族（M 循环主形态）。"""
    hits: list[dict[str, str]] = []
    for p in files:
        src = _strip_comments(_cached_text(p))
        rel = p.relative_to(ROOT).as_posix()
        for i, line in enumerate(src.split("\n"), 1):
            for m in _PERSIST_POSITIONAL.finditer(line):
                hits.append({"site": f"{rel}#L{i}", "item_id_template": m.group(1)})
    return hits


def _display_seq_sites(files: list[pathlib.Path]) -> list[str]:
    out: list[str] = []
    for p in files:
        src = _strip_comments(_cached_text(p))
        for i, line in enumerate(src.split("\n"), 1):
            if _DISPLAY_SEQ_SITE.search(line):
                out.append(f"{p.relative_to(ROOT).as_posix()}#L{i}")
    return out


def _entropy_row_key_sites(files: list[pathlib.Path]) -> list[dict[str, str]]:
    out: list[dict[str, str]] = []
    for p in files:
        src = _strip_comments(_cached_text(p))
        for i, line in enumerate(src.split("\n"), 1):
            for m in _ENTROPY_ROW_KEY.finditer(line):
                out.append({"site": f"{p.relative_to(ROOT).as_posix()}#L{i}", "prefix": m.group(1)})
    return out


def _derived_column_persist_sites(files: list[pathlib.Path]) -> list[dict[str, str]]:
    """派生列被持久化：要求「被 `calc*` 赋值」与「被 `String(row.x)` 持久化」**同时**成立。"""
    out: list[dict[str, str]] = []
    for p in files:
        src = _strip_comments(_cached_text(p))
        derived = {m.group(1) for m in _CALC_ASSIGN.finditer(src)}
        if not derived:
            continue
        rel = p.relative_to(ROOT).as_posix()
        for i, line in enumerate(src.split("\n"), 1):
            for m in _PERSIST_FIELD.finditer(line):
                if m.group(1) in derived:
                    out.append({"site": f"{rel}#L{i}", "derived_field": m.group(1)})
    return out


def _summary_persist_sites(files: list[pathlib.Path]) -> list[dict[str, str]]:
    out: list[dict[str, str]] = []
    for p in files:
        src = _strip_comments(_cached_text(p))
        rel = p.relative_to(ROOT).as_posix()
        for i, line in enumerate(src.split("\n"), 1):
            for m in _SUMMARY_ITEM.finditer(line):
                out.append({"site": f"{rel}#L{i}", "item_id": m.group(1)})
    return out


def _hardcoded_hits(files: list[pathlib.Path], name: str) -> list[str]:
    rx = _HARDCODED_PATTERNS[name]
    out: list[str] = []
    for p in files:
        src = _strip_comments(_cached_text(p))
        for i, line in enumerate(src.split("\n"), 1):
            if rx.search(line):
                out.append(f"{p.relative_to(ROOT).as_posix()}#L{i}")
    return out


# ── 宿主 dispatch 复刻 ────────────────────────────────────────────────────
_DISPATCH_BODY = re.compile(
    r"const currentSheet\s*=\s*computed\(\(\)\s*=>\s*\{(.*?)\n\}\)", re.S
)
_DISPATCH_FORMS: tuple[tuple[str, "re.Pattern[str]"], ...] = (
    ("and2", re.compile(r"if \(name\.includes\('([^']+)'\) && name\.includes\('([^']+)'\)\) return '([^']+)'")),
    ("coderx", re.compile(r"const codeMatch = name\.match\(/(M\d+-\[1-\d\])/\)")),
    ("rx_or_inc", re.compile(r"if \(name\.match\(/([^/]+)/\) \|\| name\.includes\('([^']+)'\)\) return '([^']+)'")),
    ("inc2raw", re.compile(r"if \(name\.includes\('([^']+)'\) \|\| name\.includes\('([^']+)'\)\) return name")),
    ("index", re.compile(r"if \(name\.includes\('([^']+)'\) \|\| name === '([^']+)' \|\| name === ''\) return '([^']+)'")),
    ("inc", re.compile(r"if \(name\.includes\('([^']+)'\)\) return '([^']+)'")),
)


def _dispatch_steps(host_clean: str) -> list[tuple[str, tuple[str, ...]]]:
    """从宿主源码**现读**判定步骤序列（顺序要紧 —— M10 的 `修订前` 判断排在后面而不可达）。"""
    body = _DISPATCH_BODY.search(host_clean)
    assert body, "宿主里找不到 `const currentSheet = computed(() => {…})`"
    steps: list[tuple[str, tuple[str, ...]]] = []
    for line in body.group(1).split("\n"):
        line = line.strip()
        for kind, rx in _DISPATCH_FORMS:
            m = rx.match(line)
            if m:
                steps.append((kind, m.groups()))
                break
    return steps


def _run_dispatch(steps: list[tuple[str, tuple[str, ...]]], name: str) -> str:
    for kind, g in steps:
        if kind == "and2" and g[0] in name and g[1] in name:
            return g[2]
        if kind == "coderx":
            m = re.search(g[0], name)
            if m:
                return m.group(0)
        if kind == "rx_or_inc" and (re.search(g[0], name) or g[1] in name):
            return g[2]
        if kind == "inc2raw" and (g[0] in name or g[1] in name):
            return name
        if kind == "index" and (g[0] in name or name == g[1] or name == ""):
            return g[2]
        if kind == "inc" and g[0] in name:
            return g[1]
    return name


def _dispatch_branches(host_clean: str) -> set[str]:
    return set(re.findall(r"currentSheet\s*===\s*'([^']+)'", host_clean))


_SHEET_MAP_BLOCK = re.compile(
    r"const\s+M(\d+)_SHEET_MAP\s*:\s*Record<string,\s*string>\s*=\s*\{(.*?)\n\}", re.S
)
_SHEET_MAP_PAIR = re.compile(r"^\s*'?([^':\n]+?)'?\s*:\s*'([^']*)',?\s*$", re.M)
_SHEET_MAP_FALLBACK = re.compile(r"\|\|\s*code\s*\|\|\s*'([^']+)'")


def _sheet_map_of(code: str) -> tuple[int, list[tuple[str, str]], str | None]:
    """现读某 entry 的 `M{n}_SHEET_MAP`：(声明行号, 键值对, fallback 字面量)。"""
    path = WP_COMPOSABLES / f"use{code}EntryDualMode.ts"
    src = _cached_text(path)
    block = _SHEET_MAP_BLOCK.search(src)
    assert block, f"{path.name} 里找不到 M{code[1:]}_SHEET_MAP"
    line_no = src[: block.start()].count("\n") + 1
    pairs = _SHEET_MAP_PAIR.findall(block.group(2))
    fb = _SHEET_MAP_FALLBACK.search(src)
    return line_no, pairs, (fb.group(1) if fb else None)


def _host_of(code: str) -> pathlib.Path:
    hits = sorted(
        p for p in WP_COMPONENTS.glob(f"Gt{code}*.vue")
        if re.match(rf"^Gt{code}(?![0-9])", p.name)
    )
    assert len(hits) == 1, f"Gt{code}*.vue 应恰好 1 个宿主，实得 {[p.name for p in hits]}"
    return hits[0]


def _workbook_of(code: str, manifest_slice: dict) -> pathlib.Path:
    entry = _entry_of(manifest_slice, code)
    return M_TEMPLATE_DIR / entry["template_ref"]["workbook"]


def _entry_of(manifest_slice: dict, code: str) -> dict:
    for e in manifest_slice["independent_entries"]:
        if re.match(rf"^Gt{code}(?![0-9])", e["host"]):
            return e
    raise AssertionError(f"slice 里找不到 {code} 的 entry")


def _code_of_entry(entry: dict) -> str:
    m = re.match(r"^Gt(M\d+)", entry["host"])
    assert m, entry["host"]
    return m.group(1)


def _code_of_module(rel: str) -> str:
    m = re.search(r"useM(10|[1-9])(?:Entry)?DualMode", rel)
    assert m, rel
    return "M" + m.group(1)


def _mode_gated_branches(template: str) -> list[str]:
    """模板里以 dualMode 为条件的分支指令（三要素之一：外层门控）。"""
    return re.findall(r"v-(?:if|else-if|show)=\"[^\"]*dualMode[^\"]*\"", template)


def _switch_verdict(code: str) -> tuple[str, dict[str, Any]]:
    """三要素齐验得出开关裁决（segmented + mode 门控 + OO 组件）。"""
    raw = _cached_text(_host_of(code))
    clean = _strip_comments(raw)
    template = _vue_template(clean)
    facts = {
        "segmented_sites": _lines_of(template, "el-segmented"),
        "segmented_sites_before_comment_strip": _lines_of(_vue_template(raw), "el-segmented"),
        "mode_gated_branches": sorted(set(_mode_gated_branches(template))),
        "mode_gated_branch_count": len(_mode_gated_branches(template)),
        "onlyoffice_component_sites": _lines_of(clean, "GtOnlyOfficeSheet"),
        "mounts_ac14_notice": NOTICE_COMPONENT_NAME in raw,
    }
    if facts["segmented_sites"] and facts["mode_gated_branch_count"] and facts["onlyoffice_component_sites"]:
        verdict = "switch_redeemable"
    elif facts["segmented_sites"] and not facts["mode_gated_branch_count"]:
        verdict = "switch_present_but_inert"
    else:
        verdict = "no_switch_at_all"
    return verdict, facts


def _slice_letters_on_disk() -> list[pathlib.Path]:
    """`scan_glob` 口径的全部 slice（含本轮新增的 M）。"""
    return sorted(DATA.glob("workpaper_sync_*_cycle_manifest_slice.json"))


def _load_module(name: str, path: pathlib.Path) -> ModuleType:
    spec = importlib.util.spec_from_file_location(name, path)
    assert spec and spec.loader, f"无法以 importlib 加载 {path}"
    module = importlib.util.module_from_spec(spec)
    spec.loader.exec_module(module)
    return module


def _py_function_body(source: str, signature_start: str) -> str:
    """截 Python 函数体：**先跳参数列表**再按缩进收尾（不用固定字符窗口）。"""
    i = source.find(signature_start)
    assert i >= 0, f"找不到 {signature_start!r}"
    # 跳过参数列表（可能跨多行）
    depth = 0
    j = i
    seen = False
    while j < len(source):
        ch = source[j]
        if ch == "(":
            depth += 1
            seen = True
        elif ch == ")":
            depth -= 1
            if seen and depth == 0:
                break
        j += 1
    colon = source.find(":", j)
    assert colon > 0, signature_start
    lines = source[colon + 1:].split("\n")
    body: list[str] = []
    for line in lines[1:]:
        if line.strip() and not line.startswith((" ", "\t")):
            break
        body.append(line)
    return "\n".join(body)


# ════════════════════════════════════════════════════════════════════════════
# fixtures
# ════════════════════════════════════════════════════════════════════════════
@pytest.fixture(scope="module")
def manifest_slice() -> dict:
    return _load(MANIFEST_SLICE_PATH)


@pytest.fixture(scope="module")
def deletion_plan() -> dict:
    return _load(DELETION_PLAN_PATH)


@pytest.fixture(scope="module")
def paradigm() -> dict:
    return _load(PARADIGM_PATH)


@pytest.fixture(scope="module")
def full_manifest() -> dict:
    return _load(FULL_MANIFEST_PATH)


@pytest.fixture(scope="module")
def m_files() -> list[pathlib.Path]:
    return _m_cycle_files()


@pytest.fixture(scope="module")
def m1_template(manifest_slice: dict) -> pathlib.Path:
    return _workbook_of("M1", manifest_slice)


# ════════════════════════════════════════════════════════════════════════════
# 判据零：守卫自身的反向自检（判据既不许恒红也不许恒绿）
# ════════════════════════════════════════════════════════════════════════════
class TestGuardSelfChecks:
    """先证明工具本身可用 —— 没有这一层，下面全部结论都可能是扫描器写坏了。"""

    def test_strip_comments_hides_content_but_keeps_line_numbers(self) -> None:
        src = "a\n/* el-segmented */\nb\n// el-segmented\nc"
        out = _strip_comments(src)
        assert "el-segmented" not in out
        assert out.count("\n") == src.count("\n"), "剥注释改变了行数 ⇒ 全部 #Lnn 锚点会整体上移"
        assert out.split("\n")[2] == "b" and out.split("\n")[4] == "c"

    def test_strip_comments_really_hides_a_commented_marker_in_an_m_host(self) -> None:
        """真文件上验一次：M1 宿主的注释里有 `双模式` 字样，剥完必须消失。"""
        raw = _cached_text(_host_of("M1"))
        assert "双模式切换" in raw, "前提变了：M1 宿主注释里没有该字样"
        clean = _strip_comments(raw)
        assert raw.count("\n") == clean.count("\n")
        assert clean.count("双模式切换") < raw.count("双模式切换")

    def test_inside_double_quoted_string_detector_works_both_ways(self) -> None:
        live = "import x from './useM1EntryDualMode'"
        dead = '{"snippet": "import x from \'./useM1EntryDualMode\'"}'
        assert not _inside_double_quoted_string(live, live.index("from"))
        assert _inside_double_quoted_string(dead, dead.index("from"))

    def test_resolve_spec_is_path_based_not_stem_based(self) -> None:
        importer = WP_COMPONENTS / "GtM1DividendsPayable.vue"
        got = _resolve_spec("./composables/useM1EntryDualMode", importer)
        assert got == WP_COMPOSABLES / "useM1EntryDualMode"
        assert _resolve_spec("@/components/workpaper/x", importer) == WP_COMPONENTS / "x"
        assert _resolve_spec("vue", importer) is None

    def test_vue_template_extraction_does_not_stop_at_inner_template(self) -> None:
        """M1 宿主的 `<template>` 里有多层内嵌 `<template v-if>` —— 不得被提前截断。"""
        clean = _strip_comments(_cached_text(_host_of("M1")))
        block = _vue_template(clean)
        assert block.count("<template") > 2, "顶层块被内层 template 提前截断了"
        assert "el-segmented" in block

    def test_sheet_names_are_not_stripped(self, manifest_slice: dict) -> None:
        """🔴 若守卫 strip 了 sheet 名，BP-4 的 11 处缺陷会有 6 处凭空消失。"""
        m2 = _sheet_names(_workbook_of("M2", manifest_slice))
        assert " 实收资本实质性程序表 M2A" in m2, "前导空格被 strip 了"
        m6 = _sheet_names(_workbook_of("M6", manifest_slice))
        assert "未分配利润实质性程序表 M6A " in m6, "尾随空格被 strip 了"
        m7 = _sheet_names(_workbook_of("M7", manifest_slice))
        assert " 专项储备实质性程序表 M7A " in m7, "前后空格都被 strip 了"

    def test_formula_scan_needs_data_only_false(self, m1_template: pathlib.Path) -> None:
        """🔴 反向自检：`data_only=True` 必须得到 0 —— 证明公式口径是判据不是巧合。"""
        with_formula = _formula_cell_counts(m1_template, data_only=False)
        without = _formula_cell_counts(m1_template, data_only=True)
        assert sum(with_formula.values()) > 0, "M1 册里现算不到公式格 —— 扫描器坏了"
        assert sum(without.values()) == 0, (
            "`data_only=True` 也数出了公式 ⇒ 判据失真（两边一起错就不会打红）"
        )

    def test_dispatch_replication_matches_the_source_order(self) -> None:
        """复刻器必须真的按源码顺序跑：M10 的 `修订前` 判断排在 `实质性程序表` 之后 ⇒ 不可达。"""
        clean = _strip_comments(_cached_text(_host_of("M10")))
        steps = _dispatch_steps(clean)
        kinds = [k for k, _ in steps]
        assert "coderx" in kinds and "rx_or_inc" in kinds and "inc2raw" in kinds, kinds
        assert kinds.index("rx_or_inc") < kinds.index("inc2raw"), (
            "M10 的 `修订前` 分支应排在 `实质性程序表` 之后（这正是 BP-8 的死代码根因）"
        )
        assert _run_dispatch(steps, "其他权益工具实质性程序表 Q10A (修订前)") == "procedure"
        # 反向：把 inc2raw 提到最前面，结果必须变 —— 证明复刻器真的看顺序
        reordered = [s for s in steps if s[0] == "inc2raw"] + [s for s in steps if s[0] != "inc2raw"]
        assert _run_dispatch(reordered, "其他权益工具实质性程序表 Q10A (修订前)") != "procedure"

    def test_positional_scanners_are_two_distinct_families(self) -> None:
        """🔴 两族扫描器必须互不覆盖 —— 否则「M 的主形态在持久化键上」这条差异就说不通。"""
        render_only = "  { rowKey: String(idx), name }"
        persist_only = "      { itemId: `M1-M1-1-row-${n}-name`, data: { remark: x } },"
        assert _IDENTITY_KEY.search(render_only)
        assert not _PERSIST_POSITIONAL.search(render_only)
        assert _PERSIST_POSITIONAL.search(persist_only)
        assert not any(
            _POSITIONAL_TOKEN.search(m.group(2)) or _POSITIONAL_INTERP.search(m.group(2))
            for m in _IDENTITY_KEY.finditer(persist_only)
        ), "持久化键被渲染键族误捕 ⇒ 两族计数会互相污染"

    def test_display_sequence_is_not_counted_as_a_defect(self) -> None:
        """反向判据：展示序号不是身份 —— 若它被计入，缺陷数会虚高 10。"""
        line = "      seq: idx + 1,"
        assert _DISPLAY_SEQ_SITE.search(line)
        assert not _PERSIST_POSITIONAL.search(line)

    def test_switch_verdict_discriminator_catches_an_inert_host(self) -> None:
        """🔴 M 循环 inert == 0。必须证明扫描器**能**判出 inert，否则「0」与「扫描器坏了」不可区分。"""
        assert _switch_verdict("M1")[0] == "switch_redeemable"
        assert _switch_verdict("M9")[0] == "no_switch_at_all"
        # 合成一个 L5 形态的 inert 宿主：有 segmented、v-model 绑 mode、但零 mode 门控、零 OO 组件
        inert = (
            "<template>\n<div>\n<el-segmented v-model=\"dualMode.mode\" />\n"
            "<Child v-if=\"tab === 'a'\" />\n</div>\n</template>\n"
        )
        template = _vue_template(_strip_comments(inert))
        assert _lines_of(template, "el-segmented"), "合成样本自己就不含 segmented"
        assert not _mode_gated_branches(template), "合成样本的 v-if 被误判成 mode 门控"
        assert not _lines_of(_strip_comments(inert), "GtOnlyOfficeSheet")

    def test_m_cycle_denominator_is_non_vacuous(self, m_files: list[pathlib.Path]) -> None:
        assert len(m_files) > 100, f"M 域扫描分母只有 {len(m_files)} 个文件 —— 口径写坏了"
        assert any(p.name == "useM1Adjudication.ts" for p in m_files)
        assert any(p.name == "M1TabAdjudication.vue" for p in m_files)
        assert all("__tests__" not in p.as_posix() for p in m_files)

    def test_py_function_body_skips_the_parameter_list(self) -> None:
        """截函数体禁固定字符窗口：多行签名的 `:` 会骗到第一个冒号。"""
        src = (
            "def f(\n    a: int,\n    b: dict[str, int],\n) -> list[str]:\n"
            "    return marker(a)\n\n\ndef g() -> None:\n    return None\n"
        )
        body = _py_function_body(src, "def f(")
        assert "marker(a)" in body
        assert "def g()" not in body and "return None" not in body

    def test_all_required_artifacts_exist(self) -> None:
        for p in (MANIFEST_SLICE_PATH, DELETION_PLAN_PATH, PARADIGM_PATH, FULL_MANIFEST_PATH,
                  OVERLAY_PATH, CONTRACT_DIR, TEMPLATE_INDEX, M_TEMPLATE_DIR, CHECKLIST_ROUTER,
                  REGISTRY, HTML_RENDERER_REGISTRY, SHARED_BASE, NOTICE_MODULE, NOTICE_COMPONENT,
                  CONTRACT_GUARD, COVERAGE_GUARD, CONFLICTS_PY, MERGE_PY, REMATERIALIZE_PY,
                  EXTRACT_PY, CONTRACTS_PY):
            assert p.exists(), f"缺失：{p}"
        for letter, p in SIBLING_SLICE_PATHS.items():
            assert p.is_file(), f"兄弟 slice 缺失：{letter} → {p}"


# ════════════════════════════════════════════════════════════════════════════
# 判据一：slice 范围可复算
# ════════════════════════════════════════════════════════════════════════════
class TestSliceScopeIsRecomputable:

    @staticmethod
    def _m_prefixed(full_manifest: dict) -> list[dict]:
        return [
            e for e in full_manifest["entries"]
            if any(str(p).startswith("M")
                   for p in ((e.get("wp_match") or {}).get("wp_code_patterns") or []))
        ]

    def test_selection_rule_recomputes_the_entry_set(
        self, manifest_slice: dict, full_manifest: dict
    ) -> None:
        computed = {
            e["entry_id"] for e in full_manifest["entries"]
            if e.get("document_type") == "xlsx"
            and any(str(p).startswith("M")
                    for p in ((e.get("wp_match") or {}).get("wp_code_patterns") or []))
            and e.get("independent_entry")
        }
        declared = {e["entry_id"] for e in manifest_slice["independent_entries"]}
        assert declared == computed, f"多写 {declared - computed}；漏写 {computed - declared}"

    def test_scope_counters_match_the_manifest(
        self, manifest_slice: dict, full_manifest: dict
    ) -> None:
        scope = manifest_slice["slice_scope"]
        entries = manifest_slice["independent_entries"]
        prefixed = self._m_prefixed(full_manifest)
        assert scope["independent_entry_count"] == len(entries)
        assert scope["m_prefixed_entry_total"] == len(prefixed)
        assert scope["m_prefixed_independent_total"] == len(
            [e for e in prefixed if e.get("independent_entry")]
        )
        assert scope["cycle"] == "M" and scope["document_type"] == "xlsx"

    def test_every_m_entry_is_xlsx(self, full_manifest: dict) -> None:
        prefixed = self._m_prefixed(full_manifest)
        assert prefixed, "M 前缀 entry 现算为空 —— 选择规则写坏了"
        assert {e["document_type"] for e in prefixed} == {"xlsx"}

    def test_parent_duplicate_count_is_really_zero_both_ways(
        self, manifest_slice: dict, full_manifest: dict
    ) -> None:
        ids = {e["entry_id"] for e in manifest_slice["independent_entries"]}
        children = [e for e in full_manifest["entries"] if e.get("parent_entry_id") in ids]
        assert children == [], f"现算出 parent_duplicate 子入口：{[e['entry_id'] for e in children]}"
        assert manifest_slice["slice_scope"]["parent_duplicate_count"] == 0
        assert manifest_slice["honest_adjudication_summary"]["parent_duplicate_children"] == 0
        # 反向：分母非空（manifest 里**确实存在** parent_duplicate 形态，只是不属 M）
        any_parent = [e for e in full_manifest["entries"] if e.get("parent_entry_id")]
        assert any_parent, "manifest 里一条 parent_entry_id 都没有 ⇒ 这条判据是空跑"

    def test_untriggered_conditional_sections_are_absent(self, manifest_slice: dict) -> None:
        """未触发的条件节写了就是 additive 死数据。"""
        assert "parent_duplicate_summary" not in manifest_slice
        assert "currency_variant_model" not in manifest_slice
        assert "dynamic_row_identity" in manifest_slice, "M 有动态行表，该节必须在"

    def test_there_is_no_m0_workbook_and_no_m0_entry(
        self, manifest_slice: dict, full_manifest: dict
    ) -> None:
        """🔴 与 L 的 L0 形成对照：M 没有孤儿模板册。"""
        assert not [e for e in full_manifest["entries"] if "m0" in e["entry_id"].lower()]
        assert not [r for r in _indexed_relpaths() if r.startswith("M/M0")]
        files = manifest_slice["authoritative_templates"]["files"]
        assert len(files) == len(manifest_slice["independent_entries"]), (
            "M 的模板册数应与 entry 数相等（L 是差 1，判据不得照抄）"
        )
        assert all(f["belongs_to_entry"] for f in files)
        summary = manifest_slice["honest_adjudication_summary"]
        assert summary["template_files_without_owner_entry"] == 0

    def test_host_module_edges_in_the_renderer_registry_are_real(
        self, manifest_slice: dict
    ) -> None:
        """可达性判据 = htmlRendererRegistry 的**模块边**（不是符号名 grep）。"""
        registry = _strip_comments(_cached_text(HTML_RENDERER_REGISTRY))
        for entry in manifest_slice["independent_entries"]:
            stem = entry["host"][: -len(".vue")]
            specs = [
                m.group(1) for rx in _IMPORT_FORMS for line in registry.split("\n")
                for m in rx.finditer(line)
                if not _inside_double_quoted_string(line, m.start())
            ]
            resolved = {
                r.with_suffix("").as_posix()
                for r in (_resolve_spec(s, HTML_RENDERER_REGISTRY) for s in specs) if r
            }
            want = (WP_COMPONENTS / stem).as_posix()
            assert want in resolved, f"{entry['host']} 在 registry 里没有模块边 ⇒ 不可达"

    def test_no_pilot_contract_belongs_to_the_m_cycle(self, manifest_slice: dict) -> None:
        ids = {e["entry_id"] for e in manifest_slice["independent_entries"]}
        owners: dict[str, Any] = {}
        for p in sorted(CONTRACT_DIR.glob("*.json")):
            doc = _load(p)
            owners[p.name] = (doc.get("review") or {}).get("entry_id")
        assert owners, "契约目录为空 ⇒ 这条判据是空跑"
        reviewed = {k: v for k, v in owners.items() if k != CANDIDATE_CONTRACT_FILE}
        assert set(reviewed.values()) == PILOT_CONTRACT_OWNERS, reviewed
        assert owners[CANDIDATE_CONTRACT_FILE] is None, (
            "candidate 契约的 review.entry_id 应为 null（它是反例分母，不得要求非空）"
        )
        assert _load(CONTRACT_DIR / CANDIDATE_CONTRACT_FILE)["review_status"] == "candidate"
        assert not (set(owners.values()) & ids)
        assert manifest_slice["slice_scope"]["excluded_pilot_entry_count"] == 0
        assert manifest_slice["honest_adjudication_summary"]["pilot_contract_published"] == 0

    def test_no_m_adapter_is_registered(self, manifest_slice: dict) -> None:
        registry = REGISTRY.read_text(encoding="utf-8")
        assert not re.search(r"xlsx/gt-m\d", registry), "registry 里出现了 M adapter"
        for entry in manifest_slice["independent_entries"]:
            assert entry["adapter_id"] is None
        # 反向：registry 里**确实**有 pilot adapter ⇒ 判据非空跑
        assert any(owner.split("/")[-1] in registry for owner in PILOT_CONTRACT_OWNERS), (
            "registry 里一个 pilot adapter 都找不到 ⇒ 这条判据是空跑"
        )


# ════════════════════════════════════════════════════════════════════════════
# 判据二：裁决合法性（AC 1.3 / 1.4 / 1.5 / 12.8 / 12.9 + AP-1）
# ════════════════════════════════════════════════════════════════════════════
class TestAdjudicationLegality:

    def test_every_entry_has_a_binary_html_counterpart_verdict(self, manifest_slice: dict) -> None:
        for e in manifest_slice["independent_entries"]:
            assert e["html_counterpart_verdict"] in HTML_COUNTERPART_VERDICTS, e["entry_id"]
            assert e["html_counterpart_source_refs"], e["entry_id"]

    def test_single_onlyoffice_requires_no_html_counterpart(self, manifest_slice: dict) -> None:
        for e in manifest_slice["independent_entries"]:
            if e["capability"] == "single_onlyoffice":
                assert e["html_counterpart_verdict"] == "none", e["entry_id"]

    def test_adjudication_reason_is_not_circular(self, manifest_slice: dict, paradigm: dict) -> None:
        """AP-1：不得以「本任务应交付的产物尚不存在」为裁 single 的理由。"""
        ap = next(a for a in paradigm["adjudication_criteria"]["anti_patterns"] if a["id"] == "AP-1")
        markers = ap["circular_reason_markers"]
        assert markers, "AP-1 的 marker 列表为空 ⇒ 判据空跑"
        for e in manifest_slice["independent_entries"]:
            if e["capability"] is None:
                continue
            blob = json.dumps(e["adjudication"], ensure_ascii=False)
            hit = [m for m in markers if m in blob]
            assert not hit, f"{e['entry_id']} 的裁决理由含循环论证 marker：{hit}"

    def test_capability_is_null_and_pending_fields_are_complete(
        self, manifest_slice: dict, paradigm: dict
    ) -> None:
        """SR-3 右支：null 只有在三字段齐备时才合法。"""
        schema = paradigm["slice_schema"]
        required = schema["required_pending_verdict_fields"]
        semantics = schema["pending_verdict_field_semantics"]
        pending = [e for e in manifest_slice["independent_entries"] if e["capability"] is None]
        assert len(pending) == len(manifest_slice["independent_entries"]), (
            "本 slice 应全部为待裁决态"
        )
        bp_ids = {bp["id"] for bp in manifest_slice["blocking_preconditions"]}
        for e in pending:
            for key in required:
                assert key in e, f"{e['entry_id']} 缺 {key}"
            assert semantics["capability_verdict_stage"] == "non_empty_string"
            assert isinstance(e["capability_verdict_stage"], str) and e["capability_verdict_stage"]
            assert e["capability_target"] in CAPABILITY_ENUM, e["entry_id"]
            assert isinstance(e["capability_target_blocked_by"], list)
            assert e["capability_target_blocked_by"], e["entry_id"]
            unknown = set(e["capability_target_blocked_by"]) - bp_ids
            assert not unknown, f"{e['entry_id']} 引用了不存在的 BP：{unknown}"

    def test_capability_matches_honest_capability(self, manifest_slice: dict) -> None:
        for e in manifest_slice["independent_entries"]:
            assert e["capability"] == e["adjudication"]["honest_capability"], e["entry_id"]

    def test_pending_entries_carry_no_identity(self, manifest_slice: dict) -> None:
        """SR-5 的同型约束：待裁决态不得挂 adapter / bundle / candidate / representation。"""
        for e in manifest_slice["independent_entries"]:
            for key in ("adapter_id", "authority_model", "definition_bundle",
                        "instrumentation_candidate", "published_representation",
                        "scenario_profile_id"):
                assert e[key] is None, f"{e['entry_id']}.{key} 应为 null"

    def test_adjudication_carries_both_negative_reasons(self, manifest_slice: dict) -> None:
        for e in manifest_slice["independent_entries"]:
            adj = e["adjudication"]
            assert adj["not_single_html_because"], e["entry_id"]
            assert adj["not_bidirectional_because"], e["entry_id"]

    def test_broken_or_missing_switch_is_not_read_as_a_single_html_reason(
        self, manifest_slice: dict
    ) -> None:
        """🔴 Task 54 定论：AC 12.9 的判据是「无 OO 业务价值」，开关状态不是理由。"""
        res = manifest_slice["mode_switch_resolution"]
        blob = res["why_this_is_not_a_reason_to_adjudicate_single"]
        assert "无 OO 业务价值" in blob and "AC 12.8" in blob and "AC 12.9" in blob
        for e in manifest_slice["independent_entries"]:
            because = e["adjudication"]["not_single_html_because"]
            assert "sheet" in because and "公式格" in because, e["entry_id"]
            assert "不是裁 single_html 的理由" in because or "判据（无 OO 业务价值）不成立" in because

    def test_every_blocking_precondition_is_referenced_by_someone(
        self, manifest_slice: dict
    ) -> None:
        referenced: set[str] = set()
        for e in manifest_slice["independent_entries"]:
            referenced |= set(e["capability_target_blocked_by"])
        for bp in manifest_slice["blocking_preconditions"]:
            if bp["id"] in referenced:
                continue
            # 未被 entry 引用的 BP 必须在 unverifiable_reasons 里出现
            blob = json.dumps(
                [e["evidence"]["unverifiable_reasons"] for e in manifest_slice["independent_entries"]],
                ensure_ascii=False,
            )
            assert bp["id"] in blob, f"{bp['id']} 无人引用 ⇒ 哑登记"

    def test_all_blocking_preconditions_carry_task48_fields(
        self, manifest_slice: dict, paradigm: dict
    ) -> None:
        extra = paradigm["slice_schema"]["effective_from_task_48"]["required_blocking_precondition"]
        one_of = paradigm["slice_schema"]["required_blocking_precondition_one_of"]
        for bp in manifest_slice["blocking_preconditions"]:
            for key in extra:
                assert bp.get(key), f"{bp['id']} 缺 {key}（Task 48 起必填）"
            for group in one_of:
                assert any(bp.get(k) for k in group), f"{bp['id']} 缺后果说明 {group}"
            assert bp["entries"], f"{bp['id']} 的 entries 为空 ⇒ 无对象"
            for path in bp["source_refs"]:
                # source_ref 可能带 `#L<行号>` 或 JSON pointer；只解析路径段
                assert (ROOT / path.split("#")[0]).exists(), (
                    f"{bp['id']} 的 source_ref 不存在：{path}"
                )

    def test_bp4_is_the_sheet_map_blocker(self, manifest_slice: dict) -> None:
        bp = next(b for b in manifest_slice["blocking_preconditions"] if b["id"] == "BP-4")
        audit = manifest_slice["sheet_granularity_and_router_audit"]["oo_sheet_map_audit"]
        affected = {
            _entry_of(manifest_slice, c)["entry_id"] for c in M_CODES
            if audit["per_entry"][c]["pairs_missing_from_authoritative_workbook"]
            and not audit["per_entry"][c]["module_is_orphan"]
        }
        assert set(bp["entries"]) == affected, f"BP-4 的 entries 与现算不符：{bp['entries']} vs {affected}"

    def test_bp8_is_the_sheet_granularity_blocker(self, manifest_slice: dict) -> None:
        bp = next(b for b in manifest_slice["blocking_preconditions"] if b["id"] == "BP-8")
        per = {p["wp_code"]: p for p in manifest_slice["sheet_granularity_and_router_audit"]["per_entry"]}
        collapsed = {
            _entry_of(manifest_slice, c)["entry_id"] for c in M_CODES
            if per[c]["duplicated_dispatch_codes"]
        }
        assert set(bp["entries"]) == collapsed
        assert {_code_of_entry(_entry_of(manifest_slice, c)) for c in COLLAPSE_CODES} == set(COLLAPSE_CODES)

    def test_bp11_is_the_protected_summary_blocker(self, manifest_slice: dict) -> None:
        bp = next(b for b in manifest_slice["blocking_preconditions"] if b["id"] == "BP-11")
        assert bp["blocks"] == ["step_6", "step_7"]
        assert "6.6" in json.dumps(bp, ensure_ascii=False) or "AC 6.6" in json.dumps(bp, ensure_ascii=False)
        assert set(bp["entries"]) == {e["entry_id"] for e in manifest_slice["independent_entries"]}
        assert "宁缺勿造" in json.dumps(bp["observable_consequences"], ensure_ascii=False)

    def test_bp12_is_the_unreachable_child_blocker(self, manifest_slice: dict) -> None:
        bp = next(b for b in manifest_slice["blocking_preconditions"] if b["id"] == "BP-12")
        assert len(bp["entries"]) == 1
        assert bp["entries"][0] == _entry_of(manifest_slice, "M10")["entry_id"]

    def test_manifest_mirror_divergence_is_registered_not_silently_equal(
        self, manifest_slice: dict, full_manifest: dict
    ) -> None:
        """🔴 判据 = 断言「必须不一致且已登记 BP」，不是断言相等。"""
        overlay = _load(OVERLAY_PATH)["defaults_by_component"]["GtOnlyOfficeSheet"]
        by_id = {e["entry_id"]: e for e in full_manifest["entries"]}
        bp_ids = {bp["id"] for bp in manifest_slice["blocking_preconditions"]}
        assert "BP-9" in bp_ids
        for e in manifest_slice["independent_entries"]:
            mirror = e["manifest_mirror"]
            src = by_id[e["entry_id"]]
            assert mirror["capability"] == src["capability"] == overlay["capability"]
            assert mirror["html_store"] == src["html_store"] == overlay["html_store"]
            assert mirror["capability"] != e["capability"], (
                f"{e['entry_id']}：mirror 与裁决相等了 ⇒ 说明裁决被 overlay 默认值污染"
            )
            assert "BP-9" in mirror["why_not_adopted"]
            assert "AP-3" in mirror["why_not_adopted"]

    def test_ac15_is_declared_not_applicable_and_ac14_is_the_live_one(
        self, manifest_slice: dict
    ) -> None:
        summary = manifest_slice["honest_adjudication_summary"]
        assert summary["adjudicated_as_single_html"] == 0
        assert summary["adjudicated_as_single_onlyoffice"] == 0
        assert "AC 1.5" in summary["ac_15_not_applicable_because"]
        assert "不适用" in summary["ac_15_not_applicable_because"]
        res = manifest_slice["mode_switch_resolution"]
        assert "AC 1.4" in res["live_ac_is_1_4"]
        assert "el-tooltip" in res["live_ac_is_1_4"], "AC 1.4 的 tooltip 不足条款必须写明"

    def test_no_host_template_claims_bidirectional(self, manifest_slice: dict) -> None:
        for e in manifest_slice["independent_entries"]:
            raw = _cached_text(ROOT / e["host_path"])
            assert "可双向回写" not in raw, e["host"]

    def test_ac14_notice_single_source_exists_and_is_not_duplicated(self) -> None:
        assert NOTICE_MODULE.is_file() and NOTICE_COMPONENT.is_file()
        prod, _ = _statement_edges_to(NOTICE_MODULE)
        assert len(prod) == 1, prod
        assert prod[0].split("#")[0].endswith("GtEntrySyncCapabilityNotice.vue"), prod
        comp_prod, _ = _statement_edges_to(NOTICE_COMPONENT)
        assert len(comp_prod) > 10, "提示组件的宿主消费边过少 ⇒ 它不是平台单一真源"

    def test_bp7_is_registered_because_no_m_host_mounts_the_notice(
        self, manifest_slice: dict
    ) -> None:
        mounting = [
            e["host"] for e in manifest_slice["independent_entries"]
            if NOTICE_COMPONENT_NAME in _cached_text(ROOT / e["host_path"])
        ]
        assert mounting == [], f"已有宿主挂了提示组件 ⇒ BP-7 应解除：{mounting}"
        assert manifest_slice["honest_adjudication_summary"]["entries_mounting_the_ac14_notice"] == 0
        bp = next(b for b in manifest_slice["blocking_preconditions"] if b["id"] == "BP-7")
        assert bp["status"] == "open"
        # 反向：非 M 宿主**确实**挂了 ⇒ 判据非空跑
        comp_prod, _ = _statement_edges_to(NOTICE_COMPONENT)
        assert comp_prod, "全平台没有任何宿主挂它 ⇒ 判据是空跑"

    def test_ui_toolbar_gate_anchors_are_resolvable_or_explicitly_absent(
        self, manifest_slice: dict
    ) -> None:
        for e in manifest_slice["independent_entries"]:
            gate = e["ui_toolbar_gate"]
            code = _code_of_entry(e)
            _, facts = _switch_verdict(code)
            assert gate["segmented_sites_in_gate"] == facts["segmented_sites"], code
            assert (
                gate["segmented_sites_before_comment_strip"]
                == facts["segmented_sites_before_comment_strip"]
            ), code
            if gate["anchor"] is None:
                assert gate["anchor_absent_because"], code
                assert not facts["segmented_sites"], f"{code} 有 segmented 却声明无锚点"
            else:
                assert gate["anchor_absent_because"] is None, code
                assert "el-segmented" in _line_at(gate["anchor"]), (
                    f"{code} 的锚点行不含 el-segmented：{_line_at(gate['anchor'])!r}"
                )
            assert gate["mounts_ac14_notice"] == facts["mounts_ac14_notice"]


# ════════════════════════════════════════════════════════════════════════════
# 判据三：HTML 对端 source-backed
# ════════════════════════════════════════════════════════════════════════════
class TestHtmlCounterpartIsSourceBacked:

    def test_source_refs_point_at_real_paths_and_lines(self, manifest_slice: dict) -> None:
        for e in manifest_slice["independent_entries"]:
            for ref in e["html_counterpart_source_refs"]:
                path = _resolve_repo(ref)
                assert path.is_file(), f"{e['entry_id']} 的 source_ref 不存在：{ref}"
                assert _line_at(ref) is not None

    def test_formdata_refs_really_hit_the_prefix_filter_and_the_calls(
        self, manifest_slice: dict
    ) -> None:
        """逐 entry 验四条锚点：ITEM_PREFIX 声明 / startsWith 过滤 / api.get / api.put。"""
        for e in manifest_slice["independent_entries"]:
            code = _code_of_entry(e)
            refs = [r for r in e["html_counterpart_source_refs"] if f"use{code}FormData.ts" in r]
            assert len(refs) == 4, f"{code} 的 FormData 锚点应为 4 条，实得 {refs}"
            prefix_line, filter_line, get_line, put_line = (_line_at(r) for r in refs)
            assert f"const ITEM_PREFIX = '{code}-'" in prefix_line, prefix_line
            assert "startsWith(ITEM_PREFIX)" in filter_line, filter_line
            assert "api.get(" in get_line and "checklist-responses" in get_line, get_line
            assert "api.put(" in put_line and "checklist-responses" in put_line, put_line

    def test_router_refs_are_the_get_and_put_decorators(self, manifest_slice: dict) -> None:
        for e in manifest_slice["independent_entries"]:
            refs = [r for r in e["html_counterpart_source_refs"] if "checklist_responses.py" in r]
            assert len(refs) == 2, e["entry_id"]
            get_line, put_line = (_line_at(r) for r in refs)
            assert get_line.strip().startswith("@router.get("), get_line
            assert put_line.strip().startswith("@router.put("), put_line

    def test_no_m_host_imports_the_shared_checklist_persistence(self, manifest_slice: dict) -> None:
        """🔴 K 循环 13/13 走 `useChecklistPersistence`，M 域现算 0。"""
        for e in manifest_slice["independent_entries"]:
            clean = _strip_comments(_cached_text(ROOT / e["host_path"]))
            assert "useChecklistPersistence" not in clean, e["host"]
        assert manifest_slice["honest_adjudication_summary"][
            "hosts_importing_shared_checklist_persistence"] == 0

    def test_each_entry_has_its_own_persistence_module_with_live_edges(
        self, manifest_slice: dict
    ) -> None:
        seen: set[str] = set()
        for e in manifest_slice["independent_entries"]:
            code = _code_of_entry(e)
            module = WP_COMPOSABLES / f"use{code}FormData.ts"
            assert module.is_file(), module
            prod, _ = _statement_edges_to(module)
            assert prod, f"{module.name} 没有生产边 ⇒ 它是死桩、对端结论不成立"
            assert module.as_posix() not in seen
            seen.add(module.as_posix())
        assert manifest_slice["honest_adjudication_summary"][
            "per_entry_persistence_modules"] == len(manifest_slice["independent_entries"])

    def test_item_id_prefixes_are_pairwise_non_prefixing(self, manifest_slice: dict) -> None:
        """🔴 `M1-` 与 `M10-` 的真实风险点：两两不得互为前缀。"""
        prefixes = [e["html_counterpart"]["item_id_prefix"] for e in manifest_slice["independent_entries"]]
        assert len(set(prefixes)) == len(prefixes), prefixes
        for a, b in itertools.permutations(prefixes, 2):
            assert not b.startswith(a), f"{b!r} 以 {a!r} 为前缀 ⇒ 命名空间会串"
        # 现算：前缀必须与源码里的 ITEM_PREFIX 一致
        for e in manifest_slice["independent_entries"]:
            code = _code_of_entry(e)
            src = _cached_text(WP_COMPOSABLES / f"use{code}FormData.ts")
            assert f"const ITEM_PREFIX = '{e['html_counterpart']['item_id_prefix']}'" in src

    def test_html_store_endpoint_exists_in_the_router(self, manifest_slice: dict) -> None:
        router = CHECKLIST_ROUTER.read_text(encoding="utf-8")
        assert '@router.get(""' in router and '@router.put(""' in router
        for e in manifest_slice["independent_entries"]:
            assert e["html_counterpart"]["store"] == "checklist_responses"
            assert e["html_counterpart"]["row_identity_key"] == "item_id"

    def test_template_ref_resolves_through_the_runtime_index(self, manifest_slice: dict) -> None:
        indexed = _indexed_relpaths()
        for e in manifest_slice["independent_entries"]:
            ref = e["template_ref"]
            path = M_TEMPLATE_DIR / ref["workbook"]
            assert path.is_file(), path
            assert _sha256_of(path) == ref["sha256"], ref["workbook"]
            assert len(_sheet_names(path)) == ref["sheet_count"]
            assert ref["in_runtime_index"] is True
            assert f"M/{ref['workbook']}" in indexed, ref["workbook"]

    def test_each_entry_has_its_own_workbook(self, manifest_slice: dict) -> None:
        entries = manifest_slice["independent_entries"]
        assert len({e["template_ref"]["workbook"] for e in entries}) == len(entries)
        assert len({e["template_ref"]["sha256"] for e in entries}) == len(entries)

    def test_entry_profile_fields_mirror_the_manifest(
        self, manifest_slice: dict, full_manifest: dict
    ) -> None:
        by_id = {e["entry_id"]: e for e in full_manifest["entries"]}
        for e in manifest_slice["independent_entries"]:
            src = by_id[e["entry_id"]]
            assert e["editability"] == src["editability"], e["entry_id"]
            assert e["room_model"] == src["room_model"], e["entry_id"]
            assert e["canonical_resolver"] == src["canonical_resolver"], e["entry_id"]
            assert e["migration_state"] == src["migration_state"], e["entry_id"]
            assert e["mount_count"] == len(src["mounts"]), e["entry_id"]
            assert e["wp_code_pattern"] in src["wp_match"]["wp_code_patterns"], e["entry_id"]
            assert e["host_path"] == src["host_path"], e["entry_id"]


# ════════════════════════════════════════════════════════════════════════════
# 判据四：M 循环形态差异（MD-1 ~ MD-8）
# ════════════════════════════════════════════════════════════════════════════
class TestMCycleFormDifferences:

    @staticmethod
    def _diff(manifest_slice: dict, did: str) -> dict:
        return next(d for d in manifest_slice["m_cycle_form_differences"] if d["id"] == did)

    def test_all_differences_carry_a_recompute_recipe(self, manifest_slice: dict) -> None:
        diffs = manifest_slice["m_cycle_form_differences"]
        assert len(diffs) >= 8, f"只登记了 {len(diffs)} 条形态差异"
        assert [d["id"] for d in diffs] == [f"MD-{i}" for i in range(1, len(diffs) + 1)]
        for d in diffs:
            assert d["what"] and d["why_it_matters"] and d["recompute_recipe"]
            assert d["counters"], d["id"]

    def test_md1_twin_pairing_is_exclusive_and_exhaustive(self, manifest_slice: dict) -> None:
        d = self._diff(manifest_slice, "MD-1")
        modules = _dual_mode_module_files()
        orphan, live = [], []
        for p in modules:
            prod, test = _statement_edges_to(p)
            (orphan if not prod and not test else live).append(p)
        assert len(modules) == d["counters"]["module_files_total"]
        assert len(orphan) == d["counters"]["orphan"]
        assert len(live) == d["counters"]["live"]
        assert set(orphan) & set(live) == set()
        assert set(orphan) | set(live) == set(modules)
        assert len(modules) == 2 * len(manifest_slice["independent_entries"]), (
            "M 的孪生形态要求文件数 == entry 数 × 2"
        )
        assert sum(len(_cached_text(p).splitlines()) for p in orphan) == d["counters"]["orphan_lines"]
        assert sum(len(_cached_text(p).splitlines()) for p in live) == d["counters"]["live_lines"]

    def test_md2_inert_is_zero_and_the_scanner_is_not_stuck(self, manifest_slice: dict) -> None:
        d = self._diff(manifest_slice, "MD-2")
        verdicts = {c: _switch_verdict(c)[0] for c in M_CODES}
        assert sum(1 for v in verdicts.values() if v == "switch_redeemable") == d["counters"]["switch_redeemable"]
        assert sum(1 for v in verdicts.values() if v == "switch_present_but_inert") == d["counters"]["switch_present_but_inert"]
        assert sum(1 for v in verdicts.values() if v == "no_switch_at_all") == d["counters"]["no_switch_at_all"]
        assert d["counters"]["switch_present_but_inert"] == 0, "M 循环不应有 inert"
        assert set(REDEEMABLE_CODES) == {c for c, v in verdicts.items() if v == "switch_redeemable"}
        assert set(NO_SWITCH_CODES) == {c for c, v in verdicts.items() if v == "no_switch_at_all"}

    def test_md3_sheet_map_defects_recompute_both_ways(self, manifest_slice: dict) -> None:
        d = self._diff(manifest_slice, "MD-3")
        audit = manifest_slice["sheet_granularity_and_router_audit"]["oo_sheet_map_audit"]
        total_pairs = 0
        total_bad = 0
        live_bad = 0
        fallback_ok = 0
        for code in M_CODES:
            _, pairs, fb = _sheet_map_of(code)
            real = set(_sheet_names(_workbook_of(code, manifest_slice)))
            bad = [(k, v) for k, v in pairs if v not in real]
            total_pairs += len(pairs)
            total_bad += len(bad)
            declared = audit["per_entry"][code]
            assert declared["declared_pairs"] == len(pairs), code
            assert len(declared["pairs_missing_from_authoritative_workbook"]) == len(bad), code
            assert {b["declared"] for b in declared["pairs_missing_from_authoritative_workbook"]} == {v for _, v in bad}, code
            if not declared["module_is_orphan"]:
                live_bad += len(bad)
            assert declared["fallback_literal"] == fb, code
            assert declared["fallback_hits_a_real_sheet"] == (fb in real), code
            fallback_ok += int(fb in real)
        assert total_pairs == d["counters"]["declared_pairs"]
        assert total_bad == d["counters"]["pairs_missing_in_workbook"]
        assert live_bad == d["counters"]["pairs_missing_in_live_modules"]
        assert total_bad > 0, "缺陷数为 0 ⇒ 要么已修好（该删登记）要么比对器坏了"
        # 🔴 反向判据：fallback 字面量必须**全部命中** —— 证明比对器不是恒判不存在
        assert fallback_ok == len(M_CODES) == d["counters"]["fallback_literals_hitting_real_sheet"]

    def test_md4_three_q_sheets_have_three_different_treatments(self, manifest_slice: dict) -> None:
        d = self._diff(manifest_slice, "MD-4")
        q_by_code: dict[str, list[str]] = {}
        for code in M_CODES:
            hits = [s for s in _sheet_names(_workbook_of(code, manifest_slice)) if "修订前" in s]
            if hits:
                q_by_code[code] = hits
        assert sum(len(v) for v in q_by_code.values()) == d["counters"]["pre_revision_sheets"]
        assert set(q_by_code) == {"M6", "M8", "M10"}, q_by_code
        dedicated, collapsed = 0, 0
        for code, sheets in q_by_code.items():
            clean = _strip_comments(_cached_text(_host_of(code)))
            steps = _dispatch_steps(clean)
            branches = _dispatch_branches(clean)
            for s in sheets:
                got = _run_dispatch(steps, s)
                if got == "procedure":
                    collapsed += 1
                else:
                    assert got in branches, f"{code} 的 {s!r} 派到 {got!r} 但无分支"
                    dedicated += 1
        assert dedicated == d["counters"]["handled_by_dedicated_branch"] == 1
        assert collapsed == d["counters"]["collapsed_onto_procedure"] == 2
        # M6 是那唯一正确的一条
        m6_clean = _strip_comments(_cached_text(_host_of("M6")))
        assert _run_dispatch(_dispatch_steps(m6_clean), q_by_code["M6"][0]) == "skip-q6a"

    def test_md5_templates_are_one_to_one_with_entries(self, manifest_slice: dict) -> None:
        d = self._diff(manifest_slice, "MD-5")
        disk = sorted(p.name for p in M_TEMPLATE_DIR.iterdir()
                      if p.is_file() and not p.name.startswith("~$"))
        assert len(disk) == d["counters"]["template_files"]
        assert len(manifest_slice["independent_entries"]) == d["counters"]["entries"]
        assert d["counters"]["template_files"] == d["counters"]["entries"]
        assert d["counters"]["template_files_without_owner_entry"] == 0
        assert d["counters"]["m0_entries_in_manifest"] == 0

    def test_md6_live_path_has_no_localstorage_key(self, manifest_slice: dict) -> None:
        d = self._diff(manifest_slice, "MD-6")
        base = _cached_text(SHARED_BASE)
        assert "localStorage" not in base, "共享基类出现了 localStorage ⇒ MD-6 的结论翻了"
        assert "ref<WorkpaperRenderMode>('html')" in base
        assert d["counters"]["shared_base_localstorage_hits"] == 0
        live_with_key = 0
        orphan_with_key = 0
        for p in _dual_mode_module_files():
            prod, test = _statement_edges_to(p)
            has_key = bool(re.search(r"""['"`]m\d+-dual-mode""", _cached_text(p)))
            if prod or test:
                live_with_key += int(has_key)
            else:
                orphan_with_key += int(has_key)
        assert live_with_key == d["counters"]["live_modules_with_storage_key"] == 0
        assert orphan_with_key == d["counters"]["orphan_modules_with_storage_key"]
        assert orphan_with_key > 0, "orphan 也一个键都没有 ⇒ 扫描器坏了"

    def test_md7_shared_base_edges_shrink_unlike_l(self, manifest_slice: dict) -> None:
        d = self._diff(manifest_slice, "MD-7")
        prod, test = _statement_edges_to(SHARED_BASE)
        assert len(prod) == d["counters"]["production_edges"]
        assert len(test) == d["counters"]["test_edges"]
        assert len(prod) + len(test) == d["counters"]["statement_edges"]
        m_edges = [r for r in prod + test
                   if re.search(r"/use[mM](?:10|[1-9])(?![0-9])", r)
                   or re.search(r"/Gt[mM](?:10|[1-9])(?![0-9])", r)]
        assert len(m_edges) == d["counters"]["m_cycle_edges"]
        assert d["counters"]["edges_after_rewiring"] == d["counters"]["statement_edges"] - d["counters"]["m_cycle_edges"]
        assert d["counters"]["m_cycle_edges"] > 0, "M 贡献 0 条 ⇒ 与 L 同型，MD-7 的差异不成立"
        assert d["counters"]["edges_after_rewiring"] > 0, "改线后剩 0 条 ⇒ 共享基类就该删了"
        wide = _wide_scope_edge_files("useWorkpaperEntryDualMode")
        assert len(wide) == d["counters"]["wide_token_files"]
        narrow = {r.split("#")[0] for r in prod + test}
        gap = wide - narrow
        assert gap, "宽窄口径无差 ⇒ 「排除双引号内匹配」这条口径无对象、判据变装饰"
        assert any("workpaperSyncLegacyBaseline.generated.ts" in g for g in gap), gap

    def test_md8_positional_main_form_is_the_persistence_key(
        self, manifest_slice: dict, m_files: list[pathlib.Path]
    ) -> None:
        d = self._diff(manifest_slice, "MD-8")
        render = _render_key_positional_hits(m_files)
        persist = _persist_key_positional_hits(m_files)
        seqs = _display_seq_sites(m_files)
        entropy = _entropy_row_key_sites(m_files)
        assert len(render) == d["counters"]["render_key_family"]
        assert len(persist) == d["counters"]["persistence_key_family"]
        assert len(seqs) == d["counters"]["display_sequence_sites"]
        assert len(entropy) == d["counters"]["entropy_row_key_sites"]
        assert len(persist) > len(render), (
            "持久化键族不比渲染键族多 ⇒ MD-8 的差异不成立（照抄旧扫描器就够了）"
        )
        # `n` 真是下标：逐命中文件里必须有 `const n = i + 1`
        for site in {h["site"].rsplit("#", 1)[0] for h in persist}:
            src = _strip_comments(_cached_text(ROOT / site))
            assert _N_IS_INDEX.search(src), f"{site} 的 `n` 不是下标 ⇒ 该命中应排除"
        touched = {re.match(r"(M(?:10|[1-9]))-", h["item_id_template"]).group(1) for h in persist}
        assert touched == set(M_CODES), f"持久化键族只覆盖 {sorted(touched)}"
        assert d["counters"]["entries_with_persistence_key_defect"] == len(M_CODES)


# ════════════════════════════════════════════════════════════════════════════
# 判据五：orphan / live 清册
# ════════════════════════════════════════════════════════════════════════════
class TestOrphanDualModeInventory:

    @staticmethod
    def _inv(manifest_slice: dict) -> dict:
        return manifest_slice["orphan_dual_mode_inventory"]

    def test_declared_orphans_are_unreachable_and_declared_value_matches(
        self, manifest_slice: dict
    ) -> None:
        """🔴 两侧都验（Task 54 的 M20 教训）：不只验现算 0，也验声明值 == 现算值。"""
        for mod in self._inv(manifest_slice)["modules"]:
            path = ROOT / mod["file"]
            assert path.is_file(), mod["file"]
            prod, test = _statement_edges_to(path)
            assert prod == [], f"{mod['file']} 现算出生产边 {prod} ⇒ 不是 orphan"
            assert test == [], f"{mod['file']} 现算出测试边 {test} ⇒ 不是 orphan"
            assert mod["production_consumers"] == len(prod), (
                f"{mod['file']} 声明 production_consumers={mod['production_consumers']}，现算 {len(prod)}"
            )
            assert mod["test_only_consumers"] == len(test)
            assert mod["lines"] == len(_cached_text(path).splitlines())
            assert mod["orphan_order"] == "first_order"

    def test_orphans_are_all_first_order_because_there_is_no_barrel(
        self, manifest_slice: dict
    ) -> None:
        inv = self._inv(manifest_slice)
        assert not COMPOSABLES_BARREL.exists(), "出现了 barrel ⇒ 二阶可达性判据要重算"
        assert inv["counters"]["orphan_barrels"] == 0
        assert inv["counters"]["orphan_dual_mode_second_order"] == 0
        assert inv["counters"]["orphan_dual_mode_first_order"] == inv["counters"]["orphan_dual_mode_modules"]

    def test_live_modules_have_exactly_one_edge_to_the_declared_host(
        self, manifest_slice: dict
    ) -> None:
        for mod in self._inv(manifest_slice)["live_modules"]:
            path = ROOT / mod["file"]
            prod, test = _statement_edges_to(path)
            assert mod["production_consumers"] == len(prod), mod["file"]
            assert mod["test_only_consumers"] == len(test), mod["file"]
            assert len(prod) == 1, f"{mod['file']} 的生产边应恰 1 条，实得 {prod}"
            assert re.search(r"/Gt(M\d+)[A-Za-z]*\.vue#L\d+$", prod[0]), prod[0]
            assert mod["consumer_is_host"] is True
            assert mod["consumer_is_child_tab"] is False, (
                "🔴 L 循环是子 Tab 消费，M 是宿主消费 —— 判据不得照抄"
            )
            assert mod["delegates_to_shared_base"] is True
            assert "useWorkpaperEntryDualMode" in _cached_text(path)

    def test_twenty_files_partition_into_eleven_orphans_and_nine_live(
        self, manifest_slice: dict
    ) -> None:
        inv = self._inv(manifest_slice)
        disk = {p.relative_to(ROOT).as_posix() for p in _dual_mode_module_files()}
        declared_orphan = {m["file"] for m in inv["modules"]}
        declared_live = {m["file"] for m in inv["live_modules"]}
        assert declared_orphan | declared_live == disk, (
            f"清册与磁盘不等：多 {(declared_orphan | declared_live) - disk}；漏 {disk - (declared_orphan | declared_live)}"
        )
        assert declared_orphan & declared_live == set()
        assert len(declared_orphan) == inv["counters"]["orphan_dual_mode_modules"]
        assert len(declared_live) == inv["counters"]["live_dual_mode_modules"]

    def test_the_only_orphan_entry_wrapper_is_m9(self, manifest_slice: dict) -> None:
        """🔴 M9 的两个孪生都是死桩 —— 判据不得写成「Entry 封装一定是活的」。"""
        inv = self._inv(manifest_slice)
        wrappers = [m for m in inv["modules"] if "EntryDualMode" in m["file"]]
        assert len(wrappers) == inv["counters"]["orphan_entry_wrapper_twins"] == 1
        assert _code_of_module(wrappers[0]["file"]) == "M9"
        assert wrappers[0]["is_the_entry_wrapper_twin"] is True
        # M9 的两个孪生都在 orphan 容器里
        m9 = [m["file"] for m in inv["modules"] if _code_of_module(m["file"]) == "M9"]
        assert len(m9) == 2, m9

    def test_legacy_endpoint_call_counts_recompute(self, manifest_slice: dict) -> None:
        inv = self._inv(manifest_slice)
        health = "/api/workpapers/onlyoffice/health"
        orphan_hits = 0
        for mod in inv["modules"]:
            calls = health in _cached_text(ROOT / mod["file"])
            assert (health in mod["legacy_endpoints_called"]) == calls, mod["file"]
            orphan_hits += int(calls)
        assert orphan_hits == inv["counters"]["orphan_modules_calling_legacy_health_endpoint"]
        live_hits = sum(
            1 for mod in inv["live_modules"] if health in _cached_text(ROOT / mod["file"])
        )
        assert live_hits == inv["counters"]["live_modules_calling_legacy_health_endpoint"] == 0
        assert health in _cached_text(SHARED_BASE), (
            "共享基类不调 health ⇒ 「活模块经基类间接调用」的说法不成立"
        )
        cfg = sum(
            1 for p in _dual_mode_module_files() if "onlyoffice-config" in _cached_text(p)
        )
        assert cfg == inv["counters"]["modules_calling_legacy_config_endpoint"] == 0

    def test_no_m_host_calls_the_legacy_endpoints_directly(self, manifest_slice: dict) -> None:
        for e in manifest_slice["independent_entries"]:
            clean = _strip_comments(_cached_text(ROOT / e["host_path"]))
            assert "/api/workpapers/onlyoffice/health" not in clean, e["host"]
            assert "onlyoffice-config" not in clean, e["host"]
        assert manifest_slice["honest_adjudication_summary"][
            "hosts_calling_legacy_health_endpoint_directly"] == 0

    def test_no_host_has_an_inlined_iife_dual_mode(self, manifest_slice: dict) -> None:
        """🔴 H 有 4 处、K 有 7 处宿主内联实现；M 现算 0。"""
        with_edge = 0
        for e in manifest_slice["independent_entries"]:
            clean = _strip_comments(_cached_text(ROOT / e["host_path"]))
            assert "const dualMode = (() =>" not in clean, e["host"]
            has_edge = bool(re.search(r"import \{ use(M\d+)EntryDualMode \}", clean))
            # 🔴 M9 是 no_switch_at_all ⇒ 它**不该**有载体边；其余必须有显式模块边（不得内联）
            assert has_edge == (e["dual_mode_carrier"]["kind"] != "none"), (
                f"{e['host']}：载体边（{has_edge}）与 carrier.kind"
                f"（{e['dual_mode_carrier']['kind']}）不一致"
            )
            with_edge += int(has_edge)
        assert with_edge == len(REDEEMABLE_CODES), (
            f"有显式载体边的宿主应为 {len(REDEEMABLE_CODES)} 个，实得 {with_edge}"
        )

    def test_structured_mode_values_are_only_in_orphans(self, manifest_slice: dict) -> None:
        inv = self._inv(manifest_slice)
        structured = [m for m in inv["modules"] if m["mode_enum"] and "structured" in m["mode_enum"]]
        assert len(structured) == inv["counters"]["orphan_modules_with_structured_mode_value"]
        assert structured, "一个 structured 都没有 ⇒ BP-6 应解除"
        for mod in structured:
            assert "structured" in _cached_text(ROOT / mod["file"]), mod["file"]
        for mod in inv["live_modules"]:
            assert "'structured'" not in _cached_text(ROOT / mod["file"]), mod["file"]

    def test_shared_base_counts_are_recomputed_both_ways(self, manifest_slice: dict) -> None:
        base = self._inv(manifest_slice)["shared_base_preserved"]
        prod, test = _statement_edges_to(SHARED_BASE)
        assert base["production_consumers"] == len(prod)
        assert base["test_consumers"] == len(test)
        assert base["statement_position_consumers"] == len(prod) + len(test)
        m_edges = [r for r in prod + test
                   if re.search(r"/use[mM](?:10|[1-9])(?![0-9])", r)
                   or re.search(r"/Gt[mM](?:10|[1-9])(?![0-9])", r)]
        assert sorted(base["m_cycle_consumer_refs"]) == sorted(m_edges)
        assert base["m_cycle_consumers"] == len(m_edges)
        assert base["remaining_after_this_plan"] == len(prod) + len(test) - len(m_edges)
        assert base["lines"] == len(_cached_text(SHARED_BASE).splitlines())
        wide = _wide_scope_edge_files("useWorkpaperEntryDualMode")
        assert base["wide_token_files"] == len(wide)
        assert sorted(base["wide_minus_narrow"]) == sorted(wide - {r.split("#")[0] for r in prod + test})
        assert base["persists_mode_to_localstorage"] is False
        assert base["calls_legacy_health_endpoint"] is True

    def test_orphan_summary_counts_recompute(self, manifest_slice: dict) -> None:
        inv = self._inv(manifest_slice)
        counters = inv["counters"]
        assert counters["dual_mode_module_files_total"] == len(_dual_mode_module_files())
        assert counters["orphan_dual_mode_lines_total"] == sum(m["lines"] for m in inv["modules"])
        assert counters["live_dual_mode_lines_total"] == sum(m["lines"] for m in inv["live_modules"])
        assert counters["orphan_dual_mode_modules"] + counters["live_dual_mode_modules"] == \
            counters["dual_mode_module_files_total"]
        assert counters["modules_not_scoped_to_wp_id"] == sum(
            1 for p in _dual_mode_module_files() if "wpId" not in _cached_text(p)
        )


# ════════════════════════════════════════════════════════════════════════════
# 判据六：开关裁决（AC 1.4）
# ════════════════════════════════════════════════════════════════════════════
class TestModeSwitchResolution:

    @staticmethod
    def _res(manifest_slice: dict) -> dict:
        return manifest_slice["mode_switch_resolution"]

    def test_per_entry_verdicts_are_from_a_closed_enum(self, manifest_slice: dict) -> None:
        res = self._res(manifest_slice)
        assert set(res["verdict_allowed_values"]) == set(SWITCH_VERDICTS)
        for code, verdict in res["per_entry_verdict"].items():
            assert verdict in SWITCH_VERDICTS, (code, verdict)
            assert _switch_verdict(code)[0] == verdict, code
        assert set(res["per_entry_verdict"]) == set(M_CODES)

    def test_redeemable_switch_has_all_three_elements(self, manifest_slice: dict) -> None:
        """三要素：segmented + mode 门控 v-if + OO 组件。"""
        for code in REDEEMABLE_CODES:
            verdict, facts = _switch_verdict(code)
            assert verdict == "switch_redeemable", code
            assert facts["segmented_sites"], code
            assert facts["mode_gated_branch_count"] >= 1, code
            assert facts["onlyoffice_component_sites"], code
            # 结构判据：OO 组件必须真在 mode 门控的兄弟分支下（`v-else` + resolveOoSheetName）
            template = _vue_template(_strip_comments(_cached_text(_host_of(code))))
            assert re.search(
                r"<GtOnlyOfficeSheet\s+v-else[^>]*dualMode\.resolveOoSheetName\(\)", template, re.S
            ), f"{code} 的 OO 挂载不在 mode 的 v-else 分支下 ⇒ 开关其实不可兑现"

    def test_no_switch_entry_really_has_none_of_the_three(self, manifest_slice: dict) -> None:
        for code in NO_SWITCH_CODES:
            verdict, facts = _switch_verdict(code)
            assert verdict == "no_switch_at_all", code
            assert facts["segmented_sites"] == [], code
            assert facts["mode_gated_branch_count"] == 0, code
            assert facts["onlyoffice_component_sites"], (
                f"{code} 连 OO 兜底都没有 ⇒ 那是另一种形态，登记要改"
            )
            clean = _strip_comments(_cached_text(_host_of(code)))
            assert not re.search(r"import \{ useM\d+EntryDualMode \}", clean), (
                f"{code} 竟然 import 了 dual-mode 载体 ⇒ no_switch 结论不成立"
            )

    def test_inert_count_is_zero_and_that_is_a_judgement(self, manifest_slice: dict) -> None:
        res = self._res(manifest_slice)
        assert res["counters"]["entries_with_inert_switch"] == 0
        contrast = res["contrast_with_inert"]
        assert contrast["m_cycle_inert_count"] == 0
        assert "L 循环" in contrast["why_this_key_exists"]
        assert contrast["control_group_for_zero"].startswith("M9")
        # 分母非空：三要素扫描器在真宿主上有过非零命中
        assert sum(len(_switch_verdict(c)[1]["segmented_sites"]) for c in M_CODES) > 0
        assert sum(_switch_verdict(c)[1]["mode_gated_branch_count"] for c in M_CODES) > 0

    def test_carrier_import_sites_are_real_lines(self, manifest_slice: dict) -> None:
        for e in manifest_slice["independent_entries"]:
            carrier = e["dual_mode_carrier"]
            code = _code_of_entry(e)
            if carrier["site"] is None:
                assert carrier["kind"] == "none", code
                assert carrier["module"] is None, code
                continue
            line = _line_at(carrier["site"])
            assert f"use{code}EntryDualMode" in line, (code, line)
            assert carrier["module"] == f"{WP_COMPOSABLES.relative_to(ROOT).as_posix()}/use{code}EntryDualMode.ts"
            assert carrier["kind"] == "per_entry_wrapper_over_shared_base"
            assert carrier["localStorage_prefix"] is None, "活路径不应有键（MD-6）"
            assert carrier["orphan_twin_localStorage_literal"], code

    def test_switch_counters_recompute(self, manifest_slice: dict) -> None:
        res = self._res(manifest_slice)
        verdicts = [_switch_verdict(c)[0] for c in M_CODES]
        assert res["counters"]["entries_with_redeemable_switch"] == verdicts.count("switch_redeemable")
        assert res["counters"]["entries_with_inert_switch"] == verdicts.count("switch_present_but_inert")
        assert res["counters"]["entries_with_no_switch_at_all"] == verdicts.count("no_switch_at_all")
        assert res["counters"]["hosts_with_oo_mount"] == sum(
            1 for c in M_CODES if _switch_verdict(c)[1]["onlyoffice_component_sites"]
        )
        assert res["counters"]["entries_mounting_the_ac14_notice"] == sum(
            1 for c in M_CODES if _switch_verdict(c)[1]["mounts_ac14_notice"]
        )


# ════════════════════════════════════════════════════════════════════════════
# 判据七：sheet 粒度与 OO tab 名映射
# ════════════════════════════════════════════════════════════════════════════
class TestSheetGranularityAndRouter:

    @staticmethod
    def _audit(manifest_slice: dict) -> dict:
        return manifest_slice["sheet_granularity_and_router_audit"]

    def test_sheet_coverage_recomputes_three_ways(self, manifest_slice: dict) -> None:
        audit = self._audit(manifest_slice)
        per = {p["wp_code"]: p for p in audit["per_entry"]}
        assert set(per) == set(M_CODES)
        total = covered = fell = 0
        for code in M_CODES:
            clean = _strip_comments(_cached_text(_host_of(code)))
            steps = _dispatch_steps(clean)
            branches = _dispatch_branches(clean)
            sheets = _sheet_names(_workbook_of(code, manifest_slice))
            fall = [s for s in sheets if _run_dispatch(steps, s) not in branches]
            declared = per[code]
            assert declared["authoritative_sheet_count"] == len(sheets), code
            assert declared["sheets_falling_through_to_oo"] == fall, code
            assert declared["sheets_covered_by_html_child"] == len(sheets) - len(fall), code
            assert declared["dispatch_branch_count"] == len(branches), code
            entry = _entry_of(manifest_slice, code)
            assert entry["sheet_granularity"]["authoritative_sheet_count"] == len(sheets), code
            assert entry["sheet_granularity"]["sheets_falling_through_to_oo"] == fall, code
            assert entry["template_ref"]["sheet_count"] == len(sheets), code
            total += len(sheets)
            covered += len(sheets) - len(fall)
            fell += len(fall)
        counters = audit["counters"]
        assert counters["authoritative_template_sheets_total"] == total
        assert counters["sheets_covered_by_html_children"] == covered
        assert counters["sheets_falling_through_to_oo"] == fell
        assert covered + fell == total, "覆盖面 + 兜底 ≠ 总数"

    def test_oo_fallthrough_classification_recomputes(self, manifest_slice: dict) -> None:
        audit = self._audit(manifest_slice)
        kinds: dict[str, int] = {}
        for code in M_CODES:
            clean = _strip_comments(_cached_text(_host_of(code)))
            steps = _dispatch_steps(clean)
            branches = _dispatch_branches(clean)
            for s in _sheet_names(_workbook_of(code, manifest_slice)):
                if _run_dispatch(steps, s) in branches:
                    continue
                kind = (
                    "gt_custom" if s == "GT_Custom"
                    else "pre_revision_q_sheet" if "修订前" in s
                    else "reference_only" if s.startswith("参考")
                    else "marked_deleted_in_name" if "删除" in s
                    else "unmatched_by_dispatch"
                )
                kinds[kind] = kinds.get(kind, 0) + 1
        assert kinds == audit["oo_fallthrough_classification"]["counts"], kinds
        assert kinds.get("unmatched_by_dispatch", 0) == 1, (
            "只有 M10 的国企披露 sheet 应属该类（BP-12）"
        )
        assert kinds["gt_custom"] == len(M_CODES), "每本册各一张 GT_Custom"

    def test_collapse_recomputes_and_is_exactly_three_entries(self, manifest_slice: dict) -> None:
        audit = self._audit(manifest_slice)
        per = {p["wp_code"]: p for p in audit["per_entry"]}
        collapsed_codes = []
        pairs = 0
        for code in M_CODES:
            clean = _strip_comments(_cached_text(_host_of(code)))
            steps = _dispatch_steps(clean)
            got: dict[str, list[str]] = {}
            for s in _sheet_names(_workbook_of(code, manifest_slice)):
                got.setdefault(_run_dispatch(steps, s), []).append(s)
            dups = {k: v for k, v in got.items() if len(v) > 1}
            assert per[code]["duplicated_dispatch_codes"] == dups, code
            if dups:
                collapsed_codes.append(code)
                pairs += sum(len(v) - 1 for v in dups.values())
        assert collapsed_codes == list(COLLAPSE_CODES), collapsed_codes
        assert audit["counters"]["entries_with_sheet_granularity_collapse"] == len(collapsed_codes)
        assert audit["counters"]["collapsed_sheet_pairs"] == pairs

    def test_m10_pre_revision_branch_is_unreachable_dead_code(self) -> None:
        """🔴 结构判据：那条 `if` 存在于源码里，但按顺序永不命中。"""
        clean = _strip_comments(_cached_text(_host_of("M10")))
        assert "name.includes('修订前')" in clean, "前提变了：M10 没有那条判断"
        steps = _dispatch_steps(clean)
        inc2raw = [i for i, (k, _) in enumerate(steps) if k == "inc2raw"]
        rx_or_inc = [i for i, (k, _) in enumerate(steps) if k == "rx_or_inc"]
        assert inc2raw and rx_or_inc and min(rx_or_inc) < min(inc2raw)
        q = [s for s in _sheet_names(M_TEMPLATE_DIR / "M10 其他权益工具.xlsx") if "修订前" in s]
        assert len(q) == 1
        assert _run_dispatch(steps, q[0]) == "procedure", "折叠结论不成立"

    def test_m10_soe_disclosure_child_exists_but_is_unreachable(self, manifest_slice: dict) -> None:
        """BP-12：组件存在、分支存在，但权威册上取不到任何 sheet ⇒ 结构性死代码。"""
        clean = _strip_comments(_cached_text(_host_of("M10")))
        assert "currentSheet === 'disclosure-soe'" in clean
        child = WP_COMPONENTS / "m10" / "core" / "M10TabDisclosureSoe.vue"
        assert child.is_file(), child
        steps = _dispatch_steps(clean)
        sheets = _sheet_names(_workbook_of("M10", manifest_slice))
        hit = [s for s in sheets if _run_dispatch(steps, s) == "disclosure-soe"]
        assert hit == [], f"BP-12 应解除：{hit} 命中了 disclosure-soe"
        # 反向：同一分支在别的 entry 上**有**命中 ⇒ 判据非空跑
        m9_clean = _strip_comments(_cached_text(_host_of("M9")))
        m9_hit = [s for s in _sheet_names(_workbook_of("M9", manifest_slice))
                  if _run_dispatch(_dispatch_steps(m9_clean), s) == "disclosure-soe"]
        assert m9_hit, "连 M9 的国企披露都取不到 ⇒ dispatch 复刻器坏了"

    def test_sheet_map_audit_mirrors_the_entry_level_copies(self, manifest_slice: dict) -> None:
        audit = self._audit(manifest_slice)["oo_sheet_map_audit"]
        for code in M_CODES:
            entry = _entry_of(manifest_slice, code)
            assert entry["oo_sheet_map_audit"] == audit["per_entry"][code], code
            line_no, pairs, fb = _sheet_map_of(code)
            declared = audit["per_entry"][code]
            assert declared["declaration_site"].endswith(f"#L{line_no}"), code
            assert "SHEET_MAP" in _line_at(declared["declaration_site"]), code
            assert declared["declared_pairs"] == len(pairs)
            assert declared["fallback_literal"] == fb

    def test_sheet_map_counters_recompute(self, manifest_slice: dict) -> None:
        counters = self._audit(manifest_slice)["oo_sheet_map_audit"]["counters"]
        pairs = bad = live_bad = orphan_bad = affected = fb_ok = 0
        orphan_files = {m["file"] for m in manifest_slice["orphan_dual_mode_inventory"]["modules"]}
        for code in M_CODES:
            _, ps, fb = _sheet_map_of(code)
            real = set(_sheet_names(_workbook_of(code, manifest_slice)))
            miss = [v for _, v in ps if v not in real]
            pairs += len(ps)
            bad += len(miss)
            module = f"{WP_COMPOSABLES.relative_to(ROOT).as_posix()}/use{code}EntryDualMode.ts"
            if module in orphan_files:
                orphan_bad += len(miss)
            else:
                live_bad += len(miss)
            affected += int(bool(miss))
            fb_ok += int(fb in real)
        assert counters["declared_pairs_total"] == pairs
        assert counters["pairs_missing_from_authoritative_workbook"] == bad
        assert counters["pairs_hitting_a_real_sheet"] == pairs - bad
        assert counters["pairs_missing_in_live_modules"] == live_bad
        assert counters["pairs_missing_in_orphan_modules"] == orphan_bad
        assert counters["entries_with_any_missing_pair"] == affected
        assert counters["fallback_literals_total"] == len(M_CODES)
        assert counters["fallback_literals_hitting_a_real_sheet"] == fb_ok == len(M_CODES)

    def test_verdict_declares_the_ningquewuzao_stance(self, manifest_slice: dict) -> None:
        verdict = self._audit(manifest_slice)["verdict"]
        assert "宁缺勿造" in verdict
        assert "不改宿主" in verdict and "不改模板" in verdict


# ════════════════════════════════════════════════════════════════════════════
# 判据八：Property 24 —— 公式 / 分类 summary 保持受保护（AC 6.6）
# ════════════════════════════════════════════════════════════════════════════
class TestProperty24ProtectedFormulaAndSummary:
    """tasks.md 点名项。源侧清册 + 消费侧欠账 + 承载者结构三段都是**真分母**。"""

    @staticmethod
    def _sec(manifest_slice: dict) -> dict:
        return manifest_slice["protected_formula_and_classification_summary"]

    def test_section_declares_its_ac_property_and_stance(self, manifest_slice: dict) -> None:
        sec = self._sec(manifest_slice)
        assert sec["ac_reference"] == "6.6"
        assert sec["property_reference"] == "24"
        assert "宁缺勿造" in sec["why_this_key_exists"]
        assert "不猜 `formula_mask` 区域" in sec["why_this_key_exists"]

    def test_source_side_formula_counts_recompute_per_entry(self, manifest_slice: dict) -> None:
        sec = self._sec(manifest_slice)["source_side"]
        total = sheets = 0
        for code in M_CODES:
            per_sheet = _formula_cell_counts(_workbook_of(code, manifest_slice))
            declared = sec["per_entry"][code]
            assert declared["formula_cells"] == sum(per_sheet.values()), code
            assert declared["sheets_with_formula"] == len(per_sheet), code
            for top in declared["top_sheets"]:
                assert per_sheet.get(top["sheet"]) == top["formula_cells"], (code, top)
            entry = _entry_of(manifest_slice, code)
            assert entry["template_ref"]["formula_cells"] == sum(per_sheet.values()), code
            assert entry["template_ref"]["sheets_with_formula"] == len(per_sheet), code
            total += sum(per_sheet.values())
            sheets += len(per_sheet)
        assert sec["formula_cells_total"] == total
        assert sec["sheets_with_formula_total"] == sheets
        assert total > 0 and sheets > 0, "源侧分母为空 ⇒ Property 24 的第一段变重言式"

    def test_source_side_scan_policy_is_declared_and_enforced(self, manifest_slice: dict) -> None:
        policy = manifest_slice["authoritative_templates"]["formula_scan_policy"]
        assert "data_only=False" in policy and "data_only=True" in policy
        assert "0" in policy
        # 行为判据在 TestGuardSelfChecks 里；这里核声明与实测不矛盾
        counts = _formula_cell_counts(_workbook_of("M9", manifest_slice), data_only=True)
        assert sum(counts.values()) == 0

    def test_consumer_side_two_families_recompute(
        self, manifest_slice: dict, m_files: list[pathlib.Path]
    ) -> None:
        sec = self._sec(manifest_slice)["consumer_side"]
        derived = _derived_column_persist_sites(m_files)
        summary = _summary_persist_sites(m_files)
        assert sec["derived_column_persist_sites"] == len(derived)
        assert sec["summary_persist_sites"] == len(summary)
        assert sorted(d["site"] for d in sec["derived_column_persist"]) == sorted(d["site"] for d in derived)
        assert sorted(s["item_id"] for s in sec["summary_persist"]) == sorted(s["item_id"] for s in summary)
        assert derived and summary, "消费侧分母为空 ⇒ BP-11 应解除"
        # 结构判据：每处派生列站点所在文件里必须真有 `row.<f> = calc*(...)`
        for hit in derived:
            src = _strip_comments(_cached_text(ROOT / hit["site"].rsplit("#", 1)[0]))
            assert re.search(rf"row\.{hit['derived_field']}\s*=\s*(?:calc|compute)\w+\(", src), hit

    def test_derived_column_scan_requires_both_conditions(self, m_files: list[pathlib.Path]) -> None:
        """反向自检：只满足一个条件的站点不得计入。"""
        only_persist = (
            "      { itemId: `M1-1-row-${n}-name`, data: { remark: String(row.shareholderName) } },"
        )
        assert _PERSIST_FIELD.search(only_persist)
        assert not _CALC_ASSIGN.search(only_persist)
        only_calc = "    row.endAudited = calcAuditedAmount(a, b, c)"
        assert _CALC_ASSIGN.search(only_calc)
        assert not _PERSIST_FIELD.search(only_calc)
        # 真实分母里两条件同时成立的站点数 > 0，且严格小于「只 persist」站点总数
        both = _derived_column_persist_sites(m_files)
        loose = 0
        for p in m_files:
            src = _strip_comments(_cached_text(p))
            loose += len(_PERSIST_FIELD.findall(src))
        assert 0 < len(both) < loose, (len(both), loose)

    def test_classification_summary_highlight_is_real(self, manifest_slice: dict) -> None:
        hi = self._sec(manifest_slice)["consumer_side"]["classification_summary_highlight"]
        module = ROOT / hi["module"]
        assert module.is_file(), hi["module"]
        src = _strip_comments(_cached_text(module))
        assert re.search(r"const summary\s*:\s*ComputedRef<[^>]+>\s*=\s*computed\(", src), (
            "M10 的分类 summary 不是 computed ⇒ 「派生值」的说法不成立"
        )
        for site in hi["sites"]:
            line = _line_at(site["site"])
            assert site["item_id"] in line, (site, line)
            assert "summary.value" in line, line
        assert len(hi["sites"]) == manifest_slice["honest_adjudication_summary"][
            "classification_summary_persist_sites"]
        # 权威册里真有那张分类 sheet
        sheets = _sheet_names(_workbook_of("M10", manifest_slice))
        assert hi["authoritative_sheet"] in sheets, hi["authoritative_sheet"]
        assert hi["entry_id"] == _entry_of(manifest_slice, "M10")["entry_id"]

    def test_platform_protection_policy_enum_behaves(self, manifest_slice: dict) -> None:
        """🔴 行为判据：importlib 真加载枚举并跑 `is_protected` 谓词，不是字符串存在。"""
        declared = self._sec(manifest_slice)["platform_carrier"]["protection_policy_enum"]
        # 🔴 走**包内**导入（不是按文件路径 importlib）—— conflicts.py 的 dataclass
        #    需要真实模块身份，file-path 加载会在 `_is_type` 里拿不到 `sys.modules` 条目。
        assert (ROOT / declared["file"]).is_file(), declared["file"]
        dotted = declared["file"][len("backend/"):].removesuffix(".py").replace("/", ".")
        module = importlib.import_module(dotted)
        assert module.__file__ and pathlib.Path(module.__file__) == ROOT / declared["file"], (
            f"导入到的不是声明的那个文件：{module.__file__}"
        )
        policy = getattr(module, declared["class"])
        members = [m.value for m in policy]
        assert members == declared["members"], (members, declared["members"])
        # 🔴 两侧都验：声明的受保护成员集合必须**恰好等于**由 `is_protected` 现算出的集合。
        #    只逐个断言「声明的是受保护的」会漏掉「少写一个」与「写重复」两种谎话
        #    ——「受保护单元格」这一类是独立判据，抹掉它等于该类永不被测到。
        computed_protected = [m.value for m in policy if m.is_protected]
        assert declared["protected_members"] == computed_protected, (
            f"声明 {declared['protected_members']} != 现算 {computed_protected}"
        )
        assert len(set(declared["protected_members"])) == len(declared["protected_members"]), (
            f"protected_members 有重复项：{declared['protected_members']}"
        )
        for name in declared["protected_members"]:
            assert policy(name).is_protected is True, name
        assert policy("editable").is_protected is False
        assert policy("word_only").is_protected is False

    def test_three_readonly_sources_are_distinct_in_merge(self, manifest_slice: dict) -> None:
        declared = self._sec(manifest_slice)["platform_carrier"]["three_readonly_sources"]
        src = (ROOT / declared["file"]).read_text(encoding="utf-8")
        for member in ("read_only_formula", "read_only_auto_source", "read_only_masked_cell"):
            assert f"ProtectionPolicy.{member}" in src, member
        assert "column_in_ranges" in src and "formula_mask" in src
        assert "FieldMode.formula" in src and "FieldMode.auto_source" in src

    def test_conflict_emitter_really_wraps_the_task37_function(self, manifest_slice: dict) -> None:
        """结构判据：包装函数**体内**必须真的调用被包装者（不是同名新实现）。"""
        declared = self._sec(manifest_slice)["platform_carrier"]["conflict_emitter"]
        src = (ROOT / declared["file"]).read_text(encoding="utf-8")
        body = _py_function_body(src, f"def {declared['function']}(")
        wrapped = declared["wraps"].split(".")[-1]
        assert f"{wrapped}(" in body, (
            f"{declared['function']} 的函数体里没有调用 {wrapped} ⇒ 它是第二份实现、两侧口径会漂移"
        )
        # 被包装者真存在于声明的模块里
        extract = EXTRACT_PY.read_text(encoding="utf-8")
        assert f"def {wrapped}(" in extract, wrapped
        assert f"def {declared['tamper_assertion'].split('.')[-1]}(" in extract

    def test_mask_declaration_site_exposes_both_keys(self, manifest_slice: dict) -> None:
        declared = self._sec(manifest_slice)["platform_carrier"]["mask_declaration_site"]
        src = (ROOT / declared["file"]).read_text(encoding="utf-8")
        for key in declared["keys"]:
            assert key in src, key

    def test_no_contract_and_no_mask_exists_for_this_slice(self, manifest_slice: dict) -> None:
        """空分母那侧的前提：M 侧 contract 数 0、formula_mask 声明数 0。"""
        sec = self._sec(manifest_slice)
        ids = {e["entry_id"] for e in manifest_slice["independent_entries"]}
        owners = {
            (_load(p).get("review") or {}).get("entry_id") for p in CONTRACT_DIR.glob("*.json")
        }
        assert not (owners & ids)
        assert sec["counters"]["contracts_published_for_this_slice"] == 0
        assert sec["counters"]["formula_mask_regions_declared_for_this_slice"] == 0
        # 宁缺勿造：目录里不得出现以 m{n} 命名的契约文件
        assert not [p.name for p in CONTRACT_DIR.glob("m*.json")]

    def test_protected_counters_recompute(
        self, manifest_slice: dict, m_files: list[pathlib.Path]
    ) -> None:
        counters = self._sec(manifest_slice)["counters"]
        total = sum(
            sum(_formula_cell_counts(_workbook_of(c, manifest_slice)).values()) for c in M_CODES
        )
        sheets = sum(
            len(_formula_cell_counts(_workbook_of(c, manifest_slice))) for c in M_CODES
        )
        assert counters["formula_cells_total"] == total
        assert counters["sheets_with_formula_total"] == sheets
        assert counters["entries_with_formula_sheets"] == sum(
            1 for c in M_CODES if _formula_cell_counts(_workbook_of(c, manifest_slice))
        )
        assert counters["derived_column_persist_sites"] == len(_derived_column_persist_sites(m_files))
        assert counters["summary_persist_sites"] == len(_summary_persist_sites(m_files))
        assert counters["entries_with_summary_persist"] == len(
            {s["item_id"].split("-")[0] for s in _summary_persist_sites(m_files)}
        )

    def test_verdict_does_not_overclaim(self, manifest_slice: dict) -> None:
        verdict = self._sec(manifest_slice)["verdict"]
        assert "不宣称" in verdict and "分母为空" in verdict
        den = manifest_slice["property_denominators"]["property_24"]
        assert den["not_claimed_passing"] is True
        assert den["not_claimed_passing_part"]
        assert den["carrier_for_the_empty_part"]


# ════════════════════════════════════════════════════════════════════════════
# 判据九：Property 28 —— 定义漂移 fail closed（AC 6.10）
# ════════════════════════════════════════════════════════════════════════════
class TestProperty28DefinitionDriftFailClosed:

    def test_authoritative_template_digests_recompute(self, manifest_slice: dict) -> None:
        for f in manifest_slice["authoritative_templates"]["files"]:
            path = M_TEMPLATE_DIR / f["name"]
            assert path.is_file(), f["name"]
            assert path.stat().st_size == f["size"], f["name"]
            assert _sha256_of(path) == f["sha256"], f["name"]
            assert len(_sheet_names(path)) == f["sheet_count"], f["name"]

    def test_registered_file_set_equals_the_disk_set(self, manifest_slice: dict) -> None:
        declared = {f["name"] for f in manifest_slice["authoritative_templates"]["files"]}
        disk = {p.name for p in M_TEMPLATE_DIR.iterdir()
                if p.is_file() and not p.name.startswith("~$")}
        assert declared == disk, f"多写 {declared - disk}；漏写 {disk - declared}"

    def test_lock_files_are_really_absent(self, manifest_slice: dict) -> None:
        locks = [p.name for p in M_TEMPLATE_DIR.iterdir() if p.name.startswith("~$")]
        assert locks == [], f"出现锁文件：{locks}"
        assert "~$" in manifest_slice["authoritative_templates"]["lock_file_policy"]

    def test_reference_copy_status_is_honest(self, manifest_slice: dict) -> None:
        tpl = manifest_slice["authoritative_templates"]
        assert tpl["reference_copy_status"] == "absent_in_worktree"
        assert not REFERENCE_COPY_DIR.is_dir(), (
            "参考副本目录出现了 ⇒ 声明要改成 present 并真做比对"
        )
        assert "不存在" in tpl["reference_copy_note"]

    def test_template_owner_mapping_is_a_bijection(self, manifest_slice: dict) -> None:
        files = manifest_slice["authoritative_templates"]["files"]
        ids = {e["entry_id"] for e in manifest_slice["independent_entries"]}
        owners = [f["belongs_to_entry"] for f in files]
        assert all(o is not None for o in owners), "M 不应有无主模板（与 L 的 L0 相反）"
        assert set(owners) == ids
        assert len(set(owners)) == len(owners), "同一 entry 被两本册认领"

    def test_runtime_index_membership_recomputes(self, manifest_slice: dict) -> None:
        indexed = {r for r in _indexed_relpaths() if r.startswith("M/")}
        declared = {f"M/{f['name']}" for f in manifest_slice["authoritative_templates"]["files"]}
        assert indexed == declared, f"索引与声明不等：{indexed ^ declared}"
        for f in manifest_slice["authoritative_templates"]["files"]:
            assert f["in_runtime_index"] is True

    def test_sheet_inventory_recomputes_and_spaces_are_preserved(self, manifest_slice: dict) -> None:
        """🔴 若守卫 strip 了 sheet 名，BP-4 的 6 处空格类缺陷会凭空消失。"""
        spacey = 0
        for code in M_CODES:
            names = _sheet_names(_workbook_of(code, manifest_slice))
            spacey += sum(1 for n in names if n != n.strip())
        assert spacey >= 5, f"带空格的 sheet 名现算只有 {spacey} 个 ⇒ 名字被 strip 了"
        assert " 实收资本实质性程序表 M2A" in _sheet_names(_workbook_of("M2", manifest_slice))

    def test_sheet_name_map_layer_is_the_new_drift_edge(self, manifest_slice: dict) -> None:
        """AC 6.10 在 sheet-name 层的缺口：现状**没有任何** fail closed。"""
        den = manifest_slice["property_denominators"]["property_28"]
        assert "sheet-name 映射层" in den["m_cycle_denominator"]
        bp = next(b for b in manifest_slice["blocking_preconditions"] if b["id"] == "BP-4")
        assert bp["status"] == "open"
        assert "静默错位" in json.dumps(bp, ensure_ascii=False)
        # 现算：确有缺陷（分母非空）
        bad = 0
        for code in M_CODES:
            _, pairs, _ = _sheet_map_of(code)
            real = set(_sheet_names(_workbook_of(code, manifest_slice)))
            bad += sum(1 for _, v in pairs if v not in real)
        assert bad > 0

    def test_bundle_layer_drift_is_not_claimed(self, manifest_slice: dict) -> None:
        den = manifest_slice["property_denominators"]["property_28"]
        assert den["not_claimed_passing"] is True
        assert "bundle" in den["not_claimed_passing_part"]
        for e in manifest_slice["independent_entries"]:
            assert e["definition_bundle"] is None
            assert e["published_representation"] is None
        assert manifest_slice["honest_adjudication_summary"][
            "finalized_published_representations"] == 0


# ════════════════════════════════════════════════════════════════════════════
# 判据十：Property 69 —— evidence 与计数（AC 12.10）
# ════════════════════════════════════════════════════════════════════════════
class TestProperty69EvidenceAndCounters:

    def test_unverifiable_entries_carry_non_empty_reasons(self, manifest_slice: dict) -> None:
        """SR-7。"""
        for e in manifest_slice["independent_entries"]:
            ev = e["evidence"]
            assert ev["verification_state"] == "UNVERIFIABLE", e["entry_id"]
            assert isinstance(ev["unverifiable_reasons"], list) and ev["unverifiable_reasons"]
            assert ev["sync_test_run_id"] is None and ev["required_scenario_set_digest"] is None
            assert ev["browser_case"] is None
        assert manifest_slice["honest_adjudication_summary"]["entries_left_unverifiable"] == len(
            manifest_slice["independent_entries"]
        )

    def test_no_entry_claims_verified_without_a_run(self, manifest_slice: dict) -> None:
        for e in manifest_slice["independent_entries"]:
            if e["evidence"]["verification_state"] != "UNVERIFIABLE":
                assert e["evidence"]["sync_test_run_id"], e["entry_id"]
                assert e["evidence"]["required_scenario_set_digest"], e["entry_id"]

    @pytest.mark.parametrize("bp_id", ["BP-4", "BP-8", "BP-12"])
    def test_entry_specific_bp_appears_only_in_its_own_reasons(
        self, manifest_slice: dict, bp_id: str
    ) -> None:
        """🔴 两侧都验：既验「该带的带了」，也验「不该带的没带」。"""
        bp = next(b for b in manifest_slice["blocking_preconditions"] if b["id"] == bp_id)
        expected = set(bp["entries"])
        assert expected, f"{bp_id} 的 entries 为空 ⇒ 判据空跑"
        assert expected != {e["entry_id"] for e in manifest_slice["independent_entries"]}, (
            f"{bp_id} 覆盖全部 entry ⇒ 它不是 entry 级差异，本判据不适用"
        )
        for e in manifest_slice["independent_entries"]:
            blob = "\n".join(e["evidence"]["unverifiable_reasons"])
            has = bp_id in blob
            assert has == (e["entry_id"] in expected), (
                f"{e['entry_id']} 的 reasons 与 {bp_id} 的归属不一致（has={has}）"
            )

    def test_contract_test_points_at_this_guard(self, manifest_slice: dict) -> None:
        for e in manifest_slice["independent_entries"]:
            ref = e["evidence"]["contract_test"]
            assert ref == _THIS.relative_to(ROOT).as_posix(), (e["entry_id"], ref)
            assert (ROOT / ref).is_file()

    def test_summary_counters_recompute_from_the_entries(self, manifest_slice: dict) -> None:
        summary = manifest_slice["honest_adjudication_summary"]
        entries = manifest_slice["independent_entries"]
        assert summary["total_independent"] == len(entries)
        assert summary["capability_verdict_pending"] == sum(1 for e in entries if e["capability"] is None)
        assert summary["html_counterpart_exists"] == sum(
            1 for e in entries if e["html_counterpart_verdict"] == "exists")
        assert summary["html_counterpart_none"] == sum(
            1 for e in entries if e["html_counterpart_verdict"] == "none")
        for cap in CAPABILITY_ENUM:
            assert summary[f"adjudicated_as_{cap}"] == sum(
                1 for e in entries if e["capability"] == cap), cap
        assert summary["reachable_hosts_in_cycle"] == len(entries)
        assert summary["blocking_preconditions_total"] == len(manifest_slice["blocking_preconditions"])

    def test_derived_counters_recompute_from_their_own_sections(
        self, manifest_slice: dict, m_files: list[pathlib.Path]
    ) -> None:
        """🔴 逐族穷举 —— Task 54 的 M28/M29 教训是「漏了两组摘要计数」。

        本测试按 `counting_notes` 的分组把 summary 的**每一族**都算一遍，并在末尾断言
        `counting_notes` 真的覆盖了 summary 的全部非平凡键（新增计数不登记即打红）。
        """
        summary = manifest_slice["honest_adjudication_summary"]
        entries = manifest_slice["independent_entries"]
        tpl = manifest_slice["authoritative_templates"]
        audit = manifest_slice["sheet_granularity_and_router_audit"]
        smap = audit["oo_sheet_map_audit"]
        inv = manifest_slice["orphan_dual_mode_inventory"]
        switch = manifest_slice["mode_switch_resolution"]
        dyn = manifest_slice["dynamic_row_identity"]
        prot = manifest_slice["protected_formula_and_classification_summary"]

        # 模板族
        assert summary["authoritative_template_files"] == len(tpl["files"])
        assert summary["template_files_with_owner_entry"] == sum(
            1 for f in tpl["files"] if f["belongs_to_entry"])
        assert summary["template_files_without_owner_entry"] == sum(
            1 for f in tpl["files"] if not f["belongs_to_entry"])
        assert summary["authoritative_template_sheets_total"] == sum(
            f["sheet_count"] for f in tpl["files"])
        assert summary["authoritative_template_sheets_owned"] == summary["authoritative_template_sheets_total"]
        assert summary["authoritative_formula_cells_total"] == prot["source_side"]["formula_cells_total"]
        assert summary["authoritative_sheets_with_formula"] == prot["source_side"]["sheets_with_formula_total"]

        # sheet 族
        assert summary["sheets_covered_by_html_children"] == audit["counters"]["sheets_covered_by_html_children"]
        assert summary["sheets_falling_through_to_oo"] == audit["counters"]["sheets_falling_through_to_oo"]
        kinds = audit["oo_fallthrough_classification"]["counts"]
        assert summary["oo_fallthrough_gt_custom"] == kinds.get("gt_custom", 0)
        assert summary["oo_fallthrough_reference_only"] == kinds.get("reference_only", 0)
        assert summary["oo_fallthrough_marked_deleted_in_name"] == kinds.get("marked_deleted_in_name", 0)
        assert summary["oo_fallthrough_unmatched_by_dispatch"] == kinds.get("unmatched_by_dispatch", 0)
        assert summary["oo_fallthrough_pre_revision_q_sheet"] == kinds.get("pre_revision_q_sheet", 0)
        assert sum(kinds.values()) == summary["sheets_falling_through_to_oo"]
        assert summary["entries_with_sheet_granularity_collapse"] == audit["counters"]["entries_with_sheet_granularity_collapse"]
        assert summary["collapsed_sheet_pairs"] == audit["counters"]["collapsed_sheet_pairs"]
        assert summary["pre_revision_q_sheets"] == audit["counters"]["pre_revision_q_sheets"]

        # OO sheet 名映射族
        assert summary["oo_sheet_map_declared_pairs"] == smap["counters"]["declared_pairs_total"]
        assert summary["oo_sheet_map_pairs_hitting_a_real_sheet"] == smap["counters"]["pairs_hitting_a_real_sheet"]
        assert summary["oo_sheet_map_pairs_missing_from_workbook"] == smap["counters"]["pairs_missing_from_authoritative_workbook"]
        assert summary["oo_sheet_map_pairs_missing_in_live_modules"] == smap["counters"]["pairs_missing_in_live_modules"]
        assert summary["oo_sheet_map_pairs_missing_in_orphan_modules"] == smap["counters"]["pairs_missing_in_orphan_modules"]
        assert summary["oo_sheet_map_entries_affected"] == smap["counters"]["entries_with_any_missing_pair"]
        assert summary["oo_sheet_map_fallback_literals_hitting_a_real_sheet"] == smap["counters"]["fallback_literals_hitting_a_real_sheet"]

        # dual-mode 族
        ic = inv["counters"]
        for key in ("dual_mode_module_files_total", "orphan_dual_mode_modules",
                    "orphan_dual_mode_first_order", "orphan_dual_mode_second_order",
                    "orphan_barrels", "orphan_dual_mode_lines_total",
                    "orphan_entry_wrapper_twins", "live_dual_mode_modules",
                    "live_dual_mode_lines_total",
                    "orphan_modules_calling_legacy_health_endpoint",
                    "live_modules_calling_legacy_health_endpoint",
                    "orphan_modules_with_structured_mode_value",
                    "orphan_modules_with_storage_key", "live_modules_with_storage_key"):
            assert summary[key] == ic[key], key
        assert summary["dual_mode_modules_calling_legacy_config_endpoint"] == ic["modules_calling_legacy_config_endpoint"]
        assert summary["dual_mode_modules_not_scoped_to_wp_id"] == ic["modules_not_scoped_to_wp_id"]
        base = inv["shared_base_preserved"]
        assert summary["shared_base_statement_consumers"] == base["statement_position_consumers"]
        assert summary["shared_base_m_cycle_consumers"] == base["m_cycle_consumers"]
        assert summary["shared_base_consumers_after_rewiring"] == base["remaining_after_this_plan"]
        assert summary["shared_base_wide_token_files"] == base["wide_token_files"]

        # 开关族
        for key in ("entries_with_redeemable_switch", "entries_with_inert_switch",
                    "entries_with_no_switch_at_all", "entries_mounting_the_ac14_notice",
                    "hosts_with_oo_mount"):
            assert summary[key] == switch["counters"][key], key

        # 位置化身份 + 硬编码族
        pi = dyn["positional_identity_inventory"]["counters"]
        assert summary["m_production_files_scanned"] == len(m_files)
        assert summary["dynamic_row_identity_tables"] == len(dyn["tables"])
        assert summary["positional_identity_render_key_hits"] == pi["render_key_hits"]
        assert summary["positional_identity_persistence_key_hits"] == pi["persistence_key_hits"]
        assert summary["positional_identity_defect_hits"] == pi["defect_hits_total"]
        assert summary["positional_identity_display_sequence_sites"] == pi["display_sequence_sites"]
        assert summary["positional_identity_generated_opaque_sites"] == pi["generated_opaque_sites"]
        assert summary["entries_with_positional_identity_defects"] == pi["entries_with_positional_identity_defects"]
        assert summary["entries_without_positional_identity_defects"] == pi["entries_without_positional_identity_defects"]
        assert summary["files_with_persistence_key_defect"] == pi["files_with_persistence_key_defect"]
        hc = dyn["hardcoded_pattern_scan"]["counters"]
        assert summary["hardcoded_pattern_hits_total"] == hc["hardcoded_pattern_hits_total"]
        assert summary["hardcoded_patterns_at_zero"] == hc["hardcoded_patterns_at_zero"]
        assert summary["hardcoded_patterns_total"] == hc["hardcoded_patterns_total"]

        # 受保护族
        assert summary["derived_column_persist_sites"] == prot["counters"]["derived_column_persist_sites"]
        assert summary["summary_persist_sites"] == prot["counters"]["summary_persist_sites"]
        assert summary["classification_summary_persist_sites"] == len(
            prot["consumer_side"]["classification_summary_highlight"]["sites"])
        assert summary["formula_mask_regions_declared_for_this_slice"] == prot["counters"][
            "formula_mask_regions_declared_for_this_slice"]

    def test_counting_notes_cover_every_summary_counter(self, manifest_slice: dict) -> None:
        """🔴 新增计数不登记即打红（防「加一族计数就能随便填」）。"""
        summary = manifest_slice["honest_adjudication_summary"]
        notes = summary["counting_notes"]
        listed: set[str] = set()
        for key, value in notes.items():
            if isinstance(value, list):
                listed |= set(value)
        numeric = {k for k, v in summary.items() if isinstance(v, int) and not isinstance(v, bool)}
        missing = numeric - listed
        assert not missing, f"summary 里有未登记来源的计数：{sorted(missing)}"
        stale = listed - numeric
        assert not stale, f"counting_notes 登记了不存在的计数：{sorted(stale)}"

    def test_hardcoded_scan_is_all_zero_and_non_vacuous(
        self, manifest_slice: dict, m_files: list[pathlib.Path]
    ) -> None:
        scan = manifest_slice["dynamic_row_identity"]["hardcoded_pattern_scan"]
        assert set(scan["hits"]) == set(_HARDCODED_PATTERNS), set(scan["hits"]) ^ set(_HARDCODED_PATTERNS)
        for name in _HARDCODED_PATTERNS:
            hits = _hardcoded_hits(m_files, name)
            assert hits == scan["hits"][name], (name, hits)
            assert scan["counters"][f"hardcoded_{name}"] == len(hits), name
        assert scan["counters"]["hardcoded_pattern_hits_total"] == 0
        assert scan["counters"]["hardcoded_patterns_at_zero"] == len(_HARDCODED_PATTERNS)
        # 非空跑：正则在合成样本上必须命中
        assert _HARDCODED_PATTERNS["blankRows_literal_count"].search("blankRows(p, 3)")
        assert _HARDCODED_PATTERNS["horizontal_company_column_literals"].search("'公司1'")
        assert _HARDCODED_PATTERNS["column_key_uses_label"].search("key: c.label")

    def test_slice_counters_are_all_zero_except_unadjudicated(self, manifest_slice: dict) -> None:
        counters = manifest_slice["honest_adjudication_summary"]["slice_counters"]
        pending = sum(1 for e in manifest_slice["independent_entries"] if e["capability"] is None)
        assert counters["unadjudicated"] == pending
        assert counters["fake_bidirectional_claimed_verified"] == 0
        assert counters["bidirectional_unverified"] == 0
        assert counters["stale_evidence"] == 0
        assert "不得" in counters["note"]


# ════════════════════════════════════════════════════════════════════════════
# 判据十一：Property 70 —— 跨 entry 隔离（AC 12.12）
# ════════════════════════════════════════════════════════════════════════════
class TestProperty70CrossEntryIsolation:

    def test_all_slices_are_pairwise_disjoint(self, manifest_slice: dict) -> None:
        """🔴 配对数**现算** `C(n,2)`，不写死 45（L 那轮 36，N 那轮会是 55）。"""
        paths = _slice_letters_on_disk()
        assert MANIFEST_SLICE_PATH in paths
        sets: dict[str, set[str]] = {}
        for p in paths:
            doc = _load(p)
            sets[p.name] = {
                e["entry_id"] for e in doc.get("independent_entries", []) if isinstance(e, dict)
            }
        assert len(sets) >= 10, f"scan_glob 只命中 {len(sets)} 份 slice"
        pairs = list(itertools.combinations(sorted(sets), 2))
        assert len(pairs) == len(sets) * (len(sets) - 1) // 2
        for a, b in pairs:
            overlap = sets[a] & sets[b]
            assert not overlap, f"{a} 与 {b} 的 entry 重叠：{sorted(overlap)}"
        assert sets[MANIFEST_SLICE_PATH.name], "本 slice 的 entry 集合为空 ⇒ 判据空跑"

    def test_sibling_slices_declared_match_the_disk(self, manifest_slice: dict) -> None:
        declared = manifest_slice["sibling_slices"]
        assert len(declared) == 9, f"应收全 D/E/F/G/H/I/J/K/L 九份，实得 {declared}"
        for rel in declared:
            assert (ROOT / rel).is_file(), rel
        assert set(declared) == {
            p.relative_to(ROOT).as_posix() for p in SIBLING_SLICE_PATHS.values()
        }
        assert MANIFEST_SLICE_PATH.relative_to(ROOT).as_posix() not in declared

    def test_cross_entry_isolation_assertions_are_present(self, manifest_slice: dict) -> None:
        iso = manifest_slice["cross_entry_isolation"]
        assert iso["rule"] and isinstance(iso["assertions"], list) and iso["assertions"]
        assert len(iso["assertions"]) >= 8
        assert set(iso["production_contract_files"]) == {
            p.name for p in CONTRACT_DIR.glob("*.json") if p.name != CANDIDATE_CONTRACT_FILE
        }
        assert iso["candidate_contract_files"] == [CANDIDATE_CONTRACT_FILE]
        assert "C(len(slices), 2)" in iso["pairwise_recipe"] or "C(10,2)" in iso["pairwise_recipe"]
        assert "不得写死" in iso["pairwise_recipe"]

    def test_persistence_namespaces_are_pairwise_distinct(self, manifest_slice: dict) -> None:
        prefixes = [e["html_counterpart"]["item_id_prefix"] for e in manifest_slice["independent_entries"]]
        assert len(set(prefixes)) == len(prefixes)
        keys = [tuple(e["html_counterpart"]["table_keys"]) for e in manifest_slice["independent_entries"]]
        assert len(set(keys)) == len(keys)

    def test_deletion_plan_paths_are_globally_unique_and_disjoint(
        self, deletion_plan: dict
    ) -> None:
        mine = {m["file"] for m in deletion_plan["orphan_dual_mode_to_delete"]["modules"]}
        mine |= {m["file"] for m in deletion_plan["live_dual_mode_to_rewire_then_delete"]["modules"]}
        assert mine, "删除清册为空 ⇒ 判据空跑"
        for letter, p in SIBLING_SLICE_PATHS.items():
            other_plan = p.with_name(p.name.replace("manifest_slice", "deletion_plan"))
            if not other_plan.is_file():
                continue
            doc = _load(other_plan)
            theirs: set[str] = set()
            for key in ("orphan_dual_mode_to_delete", "live_dual_mode_to_rewire_then_delete",
                        "orphan_dual_mode_to_remove", "live_dual_mode_to_delete"):
                node = doc.get(key)
                if isinstance(node, dict):
                    theirs |= {m.get("file") for m in node.get("modules", []) if m.get("file")}
            overlap = mine & theirs
            assert not overlap, f"与 {letter} 循环的删除路径重叠：{sorted(overlap)}"


# ════════════════════════════════════════════════════════════════════════════
# 判据十二：删除清册与 slice 一致
# ════════════════════════════════════════════════════════════════════════════
class TestDeletionPlanConsistency:

    def test_plan_entries_mirror_the_slice_adjudication(
        self, deletion_plan: dict, manifest_slice: dict
    ) -> None:
        by_id = {e["entry_id"]: e for e in manifest_slice["independent_entries"]}
        assert {p["entry_id"] for p in deletion_plan["entries"]} == set(by_id)
        for p in deletion_plan["entries"]:
            e = by_id[p["entry_id"]]
            assert p["capability_adjudication"] == e["capability"]
            assert p["capability_verdict_stage"] == e["capability_verdict_stage"]
            assert p["capability_target"] == e["capability_target"]
            assert p["html_counterpart_verdict"] == e["html_counterpart_verdict"]
            assert p["adjudication_reason"] == e["adjudication"]["reason"]
            assert p["html_counterpart_source_refs"] == e["html_counterpart_source_refs"]
            assert p["host_component"] == e["host_path"]
            assert p["wp_code_pattern"] == e["wp_code_pattern"]
            assert p["must_fix_before_wiring"] == e["capability_target_blocked_by"]
            assert p["switch_is_redeemable"] == e["dual_mode_carrier"]["switch_is_redeemable"]

    def test_plan_reason_is_not_circular(self, deletion_plan: dict, paradigm: dict) -> None:
        ap = next(a for a in paradigm["adjudication_criteria"]["anti_patterns"] if a["id"] == "AP-1")
        blob = json.dumps(deletion_plan, ensure_ascii=False)
        for marker in ap["circular_reason_markers"]:
            # 允许在「为什么这是循环论证」的解释里出现；这里只查 adjudication_reason 字段
            for p in deletion_plan["entries"]:
                assert marker not in p["adjudication_reason"], (p["entry_id"], marker)
        assert blob

    def test_plan_two_classes_partition_all_dual_mode_files(self, deletion_plan: dict) -> None:
        orphan = {m["file"] for m in deletion_plan["orphan_dual_mode_to_delete"]["modules"]}
        live = {m["file"] for m in deletion_plan["live_dual_mode_to_rewire_then_delete"]["modules"]}
        disk = {p.relative_to(ROOT).as_posix() for p in _dual_mode_module_files()}
        assert orphan | live == disk
        assert orphan & live == set()
        counters = deletion_plan["counters"]
        assert counters["orphan_dual_mode_to_delete"] == len(orphan)
        assert counters["live_dual_mode_to_delete"] == len(live)
        assert counters["dual_mode_module_files_total"] == len(disk)

    def test_plan_orphans_can_be_deleted_before_step_9(self, deletion_plan: dict) -> None:
        node = deletion_plan["orphan_dual_mode_to_delete"]
        assert node["can_delete_before_step_9"] is True
        assert node["ac_reference"] == "1.7"
        for mod in node["modules"]:
            prod, test = _statement_edges_to(ROOT / mod["file"])
            assert prod == [] and test == []
            assert mod["production_consumers"] == 0 and mod["test_only_consumers"] == 0
            assert mod["action"] == "delete"
            assert mod["blocking_risk"] == "none"

    def test_plan_lives_cannot_be_deleted_before_step_9_unlike_l(self, deletion_plan: dict) -> None:
        """🔴 L5~L8 是 inert 所以可先删；M 的活封装承载真开关 ⇒ 必须先改线。"""
        node = deletion_plan["live_dual_mode_to_rewire_then_delete"]
        assert node["can_delete_before_step_9"] is False
        assert "L5~L8" in node["can_delete_before_step_9_why"] or "L 循环" in node["can_delete_before_step_9_why"]
        for mod in node["modules"]:
            prod, _ = _statement_edges_to(ROOT / mod["file"])
            assert mod["production_consumers"] == prod
            assert mod["action"] == "rewire_then_delete"
            assert mod["can_delete_before_step_9"] is False

    def test_plan_has_no_inert_or_inlined_class(self, deletion_plan: dict) -> None:
        assert "inert_switch_blocks_to_remove" not in deletion_plan
        assert "host_inlined_blocks" not in deletion_plan
        counters = deletion_plan["counters"]
        assert counters["inert_switch_blocks_to_remove"] == 0
        assert counters["host_inlined_blocks_to_remove"] == 0
        assert counters["hosts_with_inlined_second_implementation"] == 0
        for p in deletion_plan["entries"]:
            assert p["inert_switch_block_to_remove"] is None
            assert p["host_inlined_block_to_remove"] is None
            assert "恒为 null" in p["inert_switch_block_note"]
            assert "恒为 null" in p["host_inlined_block_note"]

    def test_plan_shared_numbers_agree_with_the_slice(
        self, deletion_plan: dict, manifest_slice: dict
    ) -> None:
        plan_base = deletion_plan["shared_base_preserved"]
        slice_base = manifest_slice["orphan_dual_mode_inventory"]["shared_base_preserved"]
        assert plan_base["statement_position_consumers"] == slice_base["statement_position_consumers"]
        assert plan_base["m_cycle_consumers"] == slice_base["m_cycle_consumers"]
        assert plan_base["remaining_after_this_plan"] == slice_base["remaining_after_this_plan"]
        assert plan_base["non_m_consumers"] == (
            plan_base["statement_position_consumers"] - plan_base["m_cycle_consumers"]
        )
        counters = deletion_plan["counters"]
        assert counters["shared_base_consumers_before"] == plan_base["statement_position_consumers"]
        assert counters["shared_base_consumers_after"] == plan_base["remaining_after_this_plan"]
        assert counters["shared_base_consumers_after"] < counters["shared_base_consumers_before"], (
            "M 的共享基类边应收缩（与 L 的「不变」相反）"
        )

    def test_plan_health_endpoint_after_is_one_not_zero(self, deletion_plan: dict) -> None:
        """🔴 判据不得写成「删完剩 0」—— 共享基类那一处不删。"""
        counters = deletion_plan["counters"]
        health = "/api/workpapers/onlyoffice/health"
        orphan_hits = sum(
            1 for m in deletion_plan["orphan_dual_mode_to_delete"]["modules"]
            if health in _cached_text(ROOT / m["file"])
        )
        assert counters["legacy_health_endpoint_call_sites_before"] == orphan_hits + 1
        assert counters["legacy_health_endpoint_call_sites_after"] == 1
        assert health in _cached_text(SHARED_BASE)

    def test_plan_must_not_wire_to_are_the_orphan_files(self, deletion_plan: dict) -> None:
        orphan = {m["file"] for m in deletion_plan["orphan_dual_mode_to_delete"]["modules"]}
        for p in deletion_plan["entries"]:
            assert p["must_not_wire_to"], p["entry_id"]
            assert set(p["must_not_wire_to"]) <= orphan, p["entry_id"]
            assert set(p["orphan_files_to_delete_now"]) == set(p["must_not_wire_to"])
            assert p["must_not_wire_to_why"]

    def test_plan_sheet_map_fix_class_is_real(self, deletion_plan: dict, manifest_slice: dict) -> None:
        """🔴 第三类动作：修（不是删、不是改线）。"""
        node = deletion_plan["oo_sheet_map_defects_to_fix_not_delete"]
        assert node["in_scope_of_this_deletion_plan"] is False
        assert node["ac_reference"] == "6.10"
        audit = manifest_slice["sheet_granularity_and_router_audit"]["oo_sheet_map_audit"]
        declared_codes = {d["wp_code"] for d in node["defects"]}
        computed_codes = {
            c for c in M_CODES
            if audit["per_entry"][c]["pairs_missing_from_authoritative_workbook"]
        }
        assert declared_codes == computed_codes, (declared_codes, computed_codes)
        assert node["counters"]["defect_pairs_total"] == sum(
            len(d["missing"]) for d in node["defects"])
        assert node["counters"]["entries_with_defects"] == len(node["defects"])
        for d in node["defects"]:
            assert d["fix_recipe"]
            assert _resolve_repo(d["declaration_site"]).is_file()

    def test_plan_counters_recompute(self, deletion_plan: dict) -> None:
        counters = deletion_plan["counters"]
        assert counters["entries"] == len(deletion_plan["entries"])
        orphan = deletion_plan["orphan_dual_mode_to_delete"]["modules"]
        live = deletion_plan["live_dual_mode_to_rewire_then_delete"]["modules"]
        assert counters["orphan_dual_mode_lines_total"] == sum(m["lines"] for m in orphan)
        assert counters["live_dual_mode_lines_total"] == sum(m["lines"] for m in live)
        assert counters["localStorage_prefix_declaration_sites"] == sum(
            1 for m in orphan if m["storage_literal"])
        assert counters["entries_with_extra_ui_action_required"] == sum(
            1 for p in deletion_plan["entries"] if p["extra_ui_action_required"])
        assert counters["entries_with_must_fix_before_wiring"] == sum(
            1 for p in deletion_plan["entries"] if p["must_fix_before_wiring"])
        assert counters["entries_with_must_not_wire_to"] == sum(
            1 for p in deletion_plan["entries"] if p["must_not_wire_to"])
        assert counters["entries_without_any_ui_gate_anchor"] == sum(
            1 for c in M_CODES if not _switch_verdict(c)[1]["segmented_sites"])
        assert counters["entries_with_two_orphan_twins"] == len({
            _code_of_module(m["file"]) for m in orphan
            if len([g for g in orphan if _code_of_module(g["file"]) == _code_of_module(m["file"])]) > 1
        })
        assert counters["orphan_barrel_reexports_to_drop"] == 0
        assert counters["test_files_requiring_sync_change"] == 0

    def test_plan_has_no_pilot_branch(self, deletion_plan: dict) -> None:
        assert deletion_plan["pilot_already_migrated"] is None
        assert "没有" in deletion_plan["pilot_already_migrated_note"]
        for owner in PILOT_CONTRACT_OWNERS:
            assert owner in deletion_plan["pilot_already_migrated_note"]

    def test_plan_excludes_the_q_sheets_and_the_marked_deleted_sheet(
        self, deletion_plan: dict
    ) -> None:
        blob = json.dumps(deletion_plan["excluded_from_plan"], ensure_ascii=False)
        assert "修订前" in blob
        assert "针对性测试M8-5-删除" in blob, "名字带「删除」的 sheet 必须显式排除，防误删"
        assert "参考－会计规定" in blob
        assert "M0" in blob

    def test_plan_verification_recipe_mentions_the_five_checks(self, deletion_plan: dict) -> None:
        recipe = deletion_plan["why_this_plan_has_two_deletion_object_classes"]["verification_recipe"]
        for token in ("①", "②", "③", "④", "⑤"):
            assert token in recipe, token
        assert "两侧都验" in recipe

    def test_plan_declares_the_same_properties(self, deletion_plan: dict) -> None:
        assert {p["property"] for p in deletion_plan["properties_verified"]} == {24, 28, 69, 70}
        assert "6.6" in deletion_plan["requirements_covered"]
        assert "6.10" in deletion_plan["requirements_covered"]


# ════════════════════════════════════════════════════════════════════════════
# 判据十三：范式合规
# ════════════════════════════════════════════════════════════════════════════
def _validator_module() -> ModuleType:
    return _load_module("_task55_slice_schema_validator_host", CONTRACT_GUARD)


class TestParadigmCompliance:

    def test_the_slice_passes_the_paradigm_nominated_validator(self, manifest_slice: dict) -> None:
        violations = _validator_module().validate_slice_against_schema(manifest_slice)
        assert violations == [], violations

    def test_the_validator_really_catches_a_missing_pending_field(
        self, manifest_slice: dict
    ) -> None:
        broken = json.loads(json.dumps(manifest_slice, ensure_ascii=False))
        broken["independent_entries"][0].pop("capability_target_blocked_by")
        violations = _validator_module().validate_slice_against_schema(broken)
        assert any("capability_target_blocked_by" in v for v in violations), violations

    def test_the_validator_really_catches_a_counter_mismatch(self, manifest_slice: dict) -> None:
        broken = json.loads(json.dumps(manifest_slice, ensure_ascii=False))
        broken["honest_adjudication_summary"]["slice_counters"]["unadjudicated"] = 0
        violations = _validator_module().validate_slice_against_schema(broken)
        assert any("SR-9" in v for v in violations), violations

    def test_the_validator_really_catches_a_forbidden_identity_kind(
        self, manifest_slice: dict
    ) -> None:
        broken = json.loads(json.dumps(manifest_slice, ensure_ascii=False))
        broken["dynamic_row_identity"]["tables"][0]["row_identity"]["kind"] = "array_index"
        violations = _validator_module().validate_slice_against_schema(broken)
        assert any("forbidden_identity_kinds" in v for v in violations), violations

    def test_the_validator_really_catches_a_single_with_identity(
        self, manifest_slice: dict
    ) -> None:
        broken = json.loads(json.dumps(manifest_slice, ensure_ascii=False))
        entry = broken["independent_entries"][0]
        entry["capability"] = "single_onlyoffice"
        entry["adjudication"]["honest_capability"] = "single_onlyoffice"
        entry["adapter_id"] = "xlsx/gt-m1-fake"
        violations = _validator_module().validate_slice_against_schema(broken)
        assert any("SR-5" in v or "adapter_id" in v for v in violations), violations

    def test_forbidden_identity_kinds_include_the_m_specific_one(self, manifest_slice: dict) -> None:
        forbidden = manifest_slice["dynamic_row_identity"]["forbidden_identity_kinds"]
        assert "positional_persistence_key" in forbidden, (
            "M 的主缺陷形态必须进禁用清单，否则下一轮抄过去会漏"
        )
        assert "array_index" in forbidden and "display_sequence" in forbidden

    def test_dynamic_row_identity_source_refs_are_real(self, manifest_slice: dict) -> None:
        section = manifest_slice["dynamic_row_identity"]
        assert section["tables"], "动态行表清册为空 ⇒ 条件节不该出现"
        ids = {e["entry_id"] for e in manifest_slice["independent_entries"]}
        for table in section["tables"]:
            assert table["entry_id"] in ids
            identity = table["row_identity"]
            for key in ("source_ref", "persistence_source_ref"):
                ref = identity[key]
                assert ref and _resolve_repo(ref).is_file(), (table["table_key"], key, ref)
            assert "Date.now()" in _line_at(identity["source_ref"]), identity["source_ref"]
            assert "itemId" in _line_at(identity["persistence_source_ref"])
            if identity["index_binding_source_ref"]:
                assert _N_IS_INDEX.search(_line_at(identity["index_binding_source_ref"]))
            assert "BP-10" in table["registered_as"]

    def test_paradigm_refs_and_task_number_are_right(self, manifest_slice: dict) -> None:
        assert manifest_slice["task"] == "Task 55"
        assert manifest_slice["schema_version"] == "manifest-slice:v1"
        assert manifest_slice["paradigm_ref"] == PARADIGM_PATH.relative_to(ROOT).as_posix()
        assert manifest_slice["source_manifest"] == FULL_MANIFEST_PATH.relative_to(ROOT).as_posix()
        assert _validator_module()._task_number(manifest_slice) == 55

    def test_task48_extra_entry_fields_are_present(
        self, manifest_slice: dict, paradigm: dict
    ) -> None:
        extra = paradigm["slice_schema"]["effective_from_task_48"]["required_entry"]
        assert extra, "追加字段清单为空 ⇒ 判据空跑"
        for e in manifest_slice["independent_entries"]:
            for key in extra:
                assert key in e, f"{e['entry_id']} 缺 {key}"

    def test_paradigm_bytes_are_untouched_by_this_task(self, paradigm: dict) -> None:
        """🔴 本任务不改范式一个字节 —— 双向 digest 锁死。"""
        mod = _validator_module()
        registry = next(
            p for p in paradigm["paradigm_registry"]["paradigms"]
            if p["alias"] == "legacy_deletion_paradigm"
        )
        raw = mod.raw_block_digest(mod.extract_raw_json_block(mod.load_paradigm_raw(), "paradigm"))
        canonical = mod.canonical_digest(paradigm["paradigm"])
        assert raw == registry["frozen_raw_block_sha256"], "范式 `paradigm` 原始字节变了"
        assert canonical == registry["frozen_canonical_sha256"], "范式 `paradigm` 规范化字节变了"
        assert len(paradigm["paradigm"]["steps"]) == registry["step_count"] == 7

    def test_new_sections_are_declared_as_non_conflicting(self, manifest_slice: dict) -> None:
        conflict = manifest_slice["paradigm_schema_conflict"]
        assert conflict["residual_inconsistency"] is None
        declared = set(conflict["new_section_not_in_schema"]["sections"])
        schema = _load(PARADIGM_PATH)["slice_schema"]
        schema_top = set(schema["required_top_level"])
        # 条件节由 schema 自己声明（触发才必填），不算「schema 外的追加节」
        schema_top |= {cs["section"] for cs in schema.get("conditional_sections", [])}
        # `paradigm_schema_conflict` 是 D/F/G/H/I/J/K/L 八轮既定的自述键，非本轮新增
        schema_top |= {"paradigm_schema_conflict"}
        extra = {
            k for k in manifest_slice
            if k not in schema_top and not k.startswith("_")
        }
        assert declared == extra, f"追加节未全部登记：多 {declared - extra}；漏 {extra - declared}"
        assert "protected_formula_and_classification_summary" in declared
        assert conflict["new_section_not_in_schema"]["the_one_that_is_new_this_round"]
        assert "不改" in conflict["new_section_not_in_schema"]["paradigm_json_is_read_only_here"]

    def test_unjudged_slices_registry_stays_empty(self) -> None:
        cov = _load_module("_task55_coverage_guard", COVERAGE_GUARD)
        assert cov._UNJUDGED_SLICES == {}, (
            "_UNJUDGED_SLICES 非空 ⇒ 有 slice 靠 xfail 豁免了范式校验"
        )
        assert MANIFEST_SLICE_PATH in cov._contract.slice_paths(), (
            "本轮 slice 未进入覆盖面守卫的分母"
        )

    def test_requirements_and_properties_are_declared(self, manifest_slice: dict) -> None:
        props = {p["property"] for p in manifest_slice["properties_verified"]}
        assert props == {24, 28, 69, 70}
        mapping = {p["property"]: p["validates_requirement"] for p in manifest_slice["properties_verified"]}
        assert mapping == {24: "6.6", 28: "6.10", 69: "12.10", 70: "12.12"}
        for p in manifest_slice["properties_verified"]:
            key = p["see"].split(".")[-1]
            assert key in manifest_slice["property_denominators"], p
        reqs = set(manifest_slice["requirements_covered"])
        assert {"6.6", "6.10", "12.1", "12.4", "12.10", "12.11", "12.12", "14.1"} <= reqs

    def test_property_denominators_declare_what_is_not_claimed(self, manifest_slice: dict) -> None:
        den = manifest_slice["property_denominators"]
        assert "Property 24 是新出现的" in den["why_this_key_exists"]
        assert "照抄 Property 20" in den["why_this_key_exists"], (
            "必须写明「不得照抄 Property 20 的分母论证」—— 本轮 Property 集换了"
        )
        for key in ("property_24", "property_28", "property_69", "property_70"):
            node = den[key]
            assert node["what_it_asserts"] and node["how_handled"]
            assert "not_claimed_passing" in node
            if node["not_claimed_passing"]:
                assert node["not_claimed_passing_part"], key
        partial = den["property_3_and_23_partial"]
        assert partial["not_claimed_passing"] is True
        assert partial["partially_claimed"] and partial["not_claimed_parts"]

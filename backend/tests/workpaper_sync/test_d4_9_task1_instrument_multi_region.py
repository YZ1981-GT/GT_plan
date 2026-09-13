# -*- coding: utf-8 -*-
"""D4-9 Task 1 守卫：同 sheet 双受管区 instrumentation（行为级，非符号存在）。

spec: d4-9-customer-structure-bidirectional-writeback / Task 1
Requirements: 2.1, 2.2, 1.4, 8.1, 8.3

本文件的每条断言都**真跑注入 + openpyxl 复读**，判据落在 OOXML 合法性与「openpyxl 是否
同时认出两张 Table」这个结构真相上，而不是「函数存在 / 字符串存在」。
"""

from __future__ import annotations

import io
import json
import hashlib
import zipfile
from pathlib import Path

import openpyxl
import pytest

from app.services.workpaper_sync import excel_instrumentation as EI

# D4-9 权威模板（openpyxl 实测 sheet 名 + 行段，见 spec 冻结事实）。
TEMPLATE_REL = "D/D4收入底稿.xlsx"
MANAGED_SHEET = "重要客户结构分析D4-9"

_REPO_ROOT = Path(EI.__file__).resolve().parents[4]


def _sha(data: bytes) -> str:
    return hashlib.sha256(data).hexdigest()


@pytest.fixture(scope="module")
def gate(tmp_path_factory: pytest.TempPathFactory) -> EI.ExcelIdentityCarrierGate:
    """构造门对象。

    本 Task 只验证「同 sheet 双区注入的 OOXML 合法性」，与「Task 5 探针裁决是否 stale」
    正交。当前分支的 gate 基线 digest 与磁盘上的 carrier contract 漂移（上游 WIP 遗留，
    非本 spec 引入，且 refresh 脚本尚未提交），会让 `load()` 直接 stale。为让本 Task 的
    行为判据可跑，这里用**当前磁盘真实 digest** 重建一份 tmp 基线（不改任何 carrier/anchor
    的 probe_verdict，只把 stale 门的 digest 对齐到现状），再走正常 `load()`。
    这不削弱本 Task 的判据：openpyxl 是否同时认出两张 Table 与 gate 新鲜度无关。
    """
    baseline = json.loads(EI.GATE_BASELINE_PATH.read_text(encoding="utf-8"))
    for item in baseline["tier_a_runtime"]["files"]:
        p = _REPO_ROOT / item["path"]
        if p.is_file():
            item["sha256"] = _sha(p.read_bytes())
    for item in baseline["tier_a_runtime"]["probed_templates"]:
        p = _REPO_ROOT / item["path"]
        if p.is_file():
            item["sha256"] = _sha(p.read_bytes())
    for item in baseline["tier_b_evidence"]["files"]:
        p = _REPO_ROOT / item["path"]
        if p.is_file():
            item["sha256"] = _sha(p.read_bytes())
    tmp_baseline = tmp_path_factory.mktemp("gate") / "baseline.json"
    tmp_baseline.write_text(json.dumps(baseline, ensure_ascii=False), encoding="utf-8")
    return EI.ExcelIdentityCarrierGate.load(baseline_path=tmp_baseline)


@pytest.fixture(scope="module")
def template_bytes(gate: EI.ExcelIdentityCarrierGate) -> bytes:
    path = EI.TEMPLATE_AUTHORITY_ROOT / TEMPLATE_REL
    if not path.is_file():
        pytest.skip(f"D4-9 权威模板不存在: {path}")
    return path.read_bytes()


def _current_spec() -> EI.ExcelInstrumentationSpec:
    return EI.ExcelInstrumentationSpec(
        entry_id="xlsx/gt-d4-customer-structure",
        template_id="D49C",
        template_relative_path=TEMPLATE_REL,
        managed_sheet=MANAGED_SHEET,
        first_data_row=13,
        last_data_row=22,
        footer_row=23,
        managed_last_col="G",
        uuid_col="W",
        table_name="GT_D49C_ROWS",
    )


def _prior_spec() -> EI.ExcelInstrumentationSpec:
    return EI.ExcelInstrumentationSpec(
        entry_id="xlsx/gt-d4-customer-structure",
        template_id="D49P",
        template_relative_path=TEMPLATE_REL,
        managed_sheet=MANAGED_SHEET,
        first_data_row=27,
        last_data_row=36,
        footer_row=37,
        managed_last_col="G",
        uuid_col="X",
        table_name="GT_D49P_ROWS",
    )


class TestMultiRegionInjection:
    def test_openpyxl_recognizes_both_tables(
        self, template_bytes: bytes, gate: EI.ExcelIdentityCarrierGate
    ) -> None:
        """同 sheet 双区注入后，openpyxl 必须**同时**认出两张 GT_ Table。

        这正是 design §2.1.1 判死「无条件新建第二个 <tableParts>」的判据：畸形字节
        下 openpyxl 只认最后一张，第一区 identity 丢失。
        """
        result = EI.instrument_workbook_bytes_multi(
            template_bytes, [_current_spec(), _prior_spec()], gate=gate
        )
        wb = openpyxl.load_workbook(io.BytesIO(result.instrumented_bytes))
        ws = wb[MANAGED_SHEET]
        table_names = set(ws.tables.keys())
        assert "GT_D49C_ROWS" in table_names, f"本期区 Table 丢失: {table_names}"
        assert "GT_D49P_ROWS" in table_names, f"上期区 Table 丢失: {table_names}"
        # 区域边界正确（Table ref 覆盖各自行段 + UUID 列）。
        assert ws.tables["GT_D49C_ROWS"].ref == "A13:W22"
        assert ws.tables["GT_D49P_ROWS"].ref == "A27:X36"

    def test_single_tableparts_block(
        self, template_bytes: bytes, gate: EI.ExcelIdentityCarrierGate
    ) -> None:
        """产出字节里受管 sheet 只能有**一个** <tableParts> 块（OOXML schema 硬约束）。"""
        result = EI.instrument_workbook_bytes_multi(
            template_bytes, [_current_spec(), _prior_spec()], gate=gate
        )
        with zipfile.ZipFile(io.BytesIO(result.instrumented_bytes)) as zf:
            # 定位含 <tableParts> 的受管 sheet 部件（双区注入后只应有一个 sheet 带它）。
            managed = None
            for name in zf.namelist():
                if name.startswith("xl/worksheets/") and name.endswith(".xml"):
                    blob = zf.read(name).decode("utf-8", "replace")
                    if "<tableParts" in blob and "GTROW-D49C-" in blob and "GTROW-D49P-" in blob:
                        managed = blob
                        break
            assert managed is not None, "找不到含双区行 UUID + tableParts 的受管 sheet"
            assert managed.count("<tableParts") == 1, "受管 sheet 出现多个 <tableParts> 块（非法 OOXML）"
            assert 'count="2"' in managed, "<tableParts count> 未累加到 2"
            assert managed.count("<tablePart ") == 2, "应有两个 <tablePart> 子元素"

    def test_row_uuids_per_region(
        self, template_bytes: bytes, gate: EI.ExcelIdentityCarrierGate
    ) -> None:
        """两区各自持有 10 行 UUID，落在各自 UUID 列。"""
        result = EI.instrument_workbook_bytes_multi(
            template_bytes, [_current_spec(), _prior_spec()], gate=gate
        )
        cur = result.row_uuids_by_region["D49C"]
        pri = result.row_uuids_by_region["D49P"]
        assert set(cur) == set(range(13, 23)), f"本期区行 UUID 覆盖 R13-22: {sorted(cur)}"
        assert set(pri) == set(range(27, 37)), f"上期区行 UUID 覆盖 R27-36: {sorted(pri)}"
        # 两区 UUID 值互不相同（各区自带 template_id 前缀）。
        assert not (set(cur.values()) & set(pri.values())), "两区 UUID 值不得重叠"
        assert result.gt_sync_pairs["GT_MANAGED_REGION_COUNT"] == "2"

    def test_gt_sync_region_keys_are_suffixed_by_template_id(
        self, template_bytes: bytes, gate: EI.ExcelIdentityCarrierGate
    ) -> None:
        """per-region `_GT_SYNC` 键的后缀必须是 **template_id**（D49C/D49P），不是 entry_id。

        d4-9 两区**共享同一个 entry_id**（`xlsx/gt-d4-customer-structure`）。若后缀用
        entry_id，两区的 `GT_FOOTER_ROW__{entry}` 等键会互相覆盖 —— 第一区元数据被第二区
        整片盖掉，引擎消费到的只有最后一区。行为判据：两区各自的 `GT_FOOTER_ROW__D49C` /
        `..__D49P` / `GT_MANAGED_TABLE__D49C` / `..__D49P` 都存在且值正确。
        """
        result = EI.instrument_workbook_bytes_multi(
            template_bytes, [_current_spec(), _prior_spec()], gate=gate
        )
        pairs = result.gt_sync_pairs
        # 本期区（D49C）：footer 23 / table GT_D49C_ROWS / uuid 列 W。
        assert pairs["GT_FOOTER_ROW__D49C"] == "23"
        assert pairs["GT_MANAGED_TABLE__D49C"] == "GT_D49C_ROWS"
        assert pairs["GT_ROW_UUID_COLUMN__D49C"] == "W"
        # 上期区（D49P）：footer 37 / table GT_D49P_ROWS / uuid 列 X。
        assert pairs["GT_FOOTER_ROW__D49P"] == "37"
        assert pairs["GT_MANAGED_TABLE__D49P"] == "GT_D49P_ROWS"
        assert pairs["GT_ROW_UUID_COLUMN__D49P"] == "X"
        # 两区共享 entry_id ⇒ 若曾用 entry_id 作后缀，两区的 __xlsx/... 键会互相覆盖，
        # 上面六个 template_id 后缀键就无法同时正确 —— 这条断言即是那个 bug 的行为判据。
        shared_entry = _current_spec().entry_id
        assert pairs["GT_ENTRY_ID__D49C"] == shared_entry
        assert pairs["GT_ENTRY_ID__D49P"] == shared_entry


class TestSingleRegionByteIdentity:
    """既有 358 个单区工作簿：单区注入路径 `_attach_table_part` 默认分支字节必须不变。

    `_attach_table_part` 的修改只在 sheet **已有** <tableParts> 时才走新分支（追加）；
    单区注入产出的 sheet 从无 <tableParts>，故必然走原「新建独立块」路径。这里直接对
    `_attach_table_part` 的两条件做行为断言，是最直接的回归判据（不依赖 stale gate）。
    """

    def test_no_existing_block_is_original_single_block(self) -> None:
        sheet = '<worksheet xmlns:r="ns"><sheetData/></worksheet>'
        out = EI._attach_table_part(sheet)
        assert out.count("<tableParts") == 1
        assert 'count="1"' in out
        assert f'<tablePart r:id="{EI._GT_TABLE_REL_ID}"/>' in out
        # 原实现在无尾部元素时插到 </worksheet> 前 —— 位置不变。
        assert out.endswith("</tableParts></worksheet>")

    def test_existing_block_appends_and_bumps_count(self) -> None:
        sheet = (
            '<worksheet xmlns:r="ns"><sheetData/>'
            '<tableParts count="1"><tablePart r:id="rIdGTTBL1"/></tableParts>'
            "</worksheet>"
        )
        out = EI._attach_table_part(sheet, rel_id="rIdGTTBL12")
        assert out.count("<tableParts") == 1, "不得新建第二个 <tableParts> 块"
        assert 'count="2"' in out
        assert out.count("<tablePart ") == 2


class TestStructuralFailClosed:
    def test_different_sheets_rejected(
        self, template_bytes: bytes, gate: EI.ExcelIdentityCarrierGate
    ) -> None:
        other = EI.ExcelInstrumentationSpec(
            entry_id="x", template_id="OTHER", template_relative_path=TEMPLATE_REL,
            managed_sheet="重要客户销售价格分析D4-10", first_data_row=27, last_data_row=36,
            footer_row=37, managed_last_col="G", uuid_col="X", table_name="GT_OTHER",
        )
        with pytest.raises(EI.InstrumentationError, match="同一 sheet"):
            EI.instrument_workbook_bytes_multi(template_bytes, [_current_spec(), other], gate=gate)

    def test_overlapping_rows_rejected(
        self, template_bytes: bytes, gate: EI.ExcelIdentityCarrierGate
    ) -> None:
        overlap = EI.ExcelInstrumentationSpec(
            entry_id="x", template_id="D49P", template_relative_path=TEMPLATE_REL,
            managed_sheet=MANAGED_SHEET, first_data_row=20, last_data_row=30,
            footer_row=37, managed_last_col="G", uuid_col="X", table_name="GT_D49P_ROWS",
        )
        with pytest.raises(EI.InstrumentationError, match="行段重叠"):
            EI.instrument_workbook_bytes_multi(template_bytes, [_current_spec(), overlap], gate=gate)

    def test_duplicate_uuid_column_rejected(
        self, template_bytes: bytes, gate: EI.ExcelIdentityCarrierGate
    ) -> None:
        dup_col = EI.ExcelInstrumentationSpec(
            entry_id="x", template_id="D49P", template_relative_path=TEMPLATE_REL,
            managed_sheet=MANAGED_SHEET, first_data_row=27, last_data_row=36,
            footer_row=37, managed_last_col="G", uuid_col="W", table_name="GT_D49P_ROWS",
        )
        with pytest.raises(EI.InstrumentationError, match="UUID 列"):
            EI.instrument_workbook_bytes_multi(template_bytes, [_current_spec(), dup_col], gate=gate)

    def test_duplicate_table_name_rejected(
        self, template_bytes: bytes, gate: EI.ExcelIdentityCarrierGate
    ) -> None:
        dup_tbl = EI.ExcelInstrumentationSpec(
            entry_id="x", template_id="D49P", template_relative_path=TEMPLATE_REL,
            managed_sheet=MANAGED_SHEET, first_data_row=27, last_data_row=36,
            footer_row=37, managed_last_col="G", uuid_col="X", table_name="GT_D49C_ROWS",
        )
        with pytest.raises(EI.InstrumentationError, match="displayName"):
            EI.instrument_workbook_bytes_multi(template_bytes, [_current_spec(), dup_tbl], gate=gate)

"""骨架槽位复用判据（spec d2-skeleton-slot-reuse-and-writeback-speed · design §9）。

三层：

1. 纯函数 —— 候选（模板原样 / 用户填过 / 重复身份 / 未启用）与复用计划（全空模板 / 超出槽位 /
   部分已占 / 存量布局 / projection 声明了骨架行）；
2. 身份保留门 —— 只豁免「声明退役 ∧ 骨架身份形态」；
3. 真链路 —— 新 instrument 的 D2 权威模板 + store 行走生产同一组函数（首版发布用的
   ``stage_instrumented_substrate`` / ``_identity_binding`` / ``_sibling_identity_bindings`` /
   overlay / ``adapter.materialize`` / ``adapter.extract`` / roundtrip / 未管理区域），断言
   业务数据从第 13 行开始。
"""
from __future__ import annotations

import hashlib
import io
import json
import os
import sys
import uuid
from pathlib import Path
from types import SimpleNamespace
from typing import Any

import pytest

_BACKEND = Path(__file__).resolve().parents[2]
if str(_BACKEND) not in sys.path:  # pragma: no cover
    sys.path.insert(0, str(_BACKEND))
os.environ.setdefault("DB_DISABLE_SSL", "True")

from app.services.workpaper_sync import excel_extract as X  # noqa: E402
from app.services.workpaper_sync import excel_materialize as M  # noqa: E402
from app.services.workpaper_sync import pilot_d2_large_json as P  # noqa: E402
from app.services.workpaper_sync import projection_first_publication as F  # noqa: E402
from app.services.workpaper_sync import skeleton_slot_reuse as S  # noqa: E402
from app.services.workpaper_sync.adapters import excel as AX  # noqa: E402
from app.services.workpaper_sync.adapters.base import (  # noqa: E402
    FieldValue,
    Projection,
    SubstrateRole,
)
from app.services.workpaper_sync.content_mutation import (  # noqa: E402
    assert_roundtrip_equivalent,
)
from app.services.workpaper_sync.contracts import FieldMode, ValueType  # noqa: E402
from app.services.workpaper_sync.excel_entry_gate import (  # noqa: E402
    AdapterBuild,
    FrozenEntryDefinitions,
)
from app.services.workpaper_sync.models import (  # noqa: E402
    ArtifactKind,
    ArtifactState,
    AuthorityModel,
    BundleSlot,
    BundleSlotSpec,
    DefinitionState,
)
from app.services.workpaper_sync.resolution import DefinitionBundleSnapshot  # noqa: E402
from app.services.workpaper_sync.template_skeleton_baseline import (  # noqa: E402
    SkeletonBaseline,
)

_T = P.ROWS_TABLE_KEY
_SK = [f"GTROW-D22-{row:04d}" for row in range(13, 25)]


@pytest.fixture(autouse=True)
def _fresh_declaration_cache():
    S.clear_reuse_declaration_cache()
    yield
    S.clear_reuse_declaration_cache()


# ═══════════════════════════════════════════════════════════════════════════
# 1. 复用计划（纯函数）
# ═══════════════════════════════════════════════════════════════════════════


def _fresh_physical() -> dict[int, str]:
    return {13 + offset: identity for offset, identity in enumerate(_SK)}


def test_fresh_template_three_rows_take_the_first_three_slots() -> None:
    physical = _fresh_physical()
    adoption = S.plan_slot_adoption(
        candidates=physical, physical=physical, wanted=["r1", "r2", "r3"],
        orphan=["r1", "r2", "r3"],
    )
    assert adoption.adopted == {"r1": 13, "r2": 14, "r3": 15}
    assert adoption.retired_identities == tuple(_SK[:3])


def test_more_rows_than_slots_leave_the_remainder_for_insertion() -> None:
    physical = _fresh_physical()
    orphan = [f"r{i}" for i in range(20)]
    adoption = S.plan_slot_adoption(
        candidates=physical, physical=physical, wanted=orphan, orphan=orphan
    )
    assert list(adoption.adopted) == orphan[:12], "声明序与槽位升序一一对应"
    assert sorted(adoption.adopted.values()) == list(range(13, 25))


def test_new_row_after_existing_data_takes_the_next_free_slot() -> None:
    physical = _fresh_physical()
    physical.update({13: "r1", 14: "r2", 15: "r3"})
    candidates = {row: ident for row, ident in physical.items() if ident.startswith("GTROW")}
    adoption = S.plan_slot_adoption(
        candidates=candidates, physical=physical, wanted=["r1", "r2", "r3", "r4"],
        orphan=["r4"],
    )
    assert adoption.adopted == {"r4": 16}
    assert adoption.retired == {16: "GTROW-D22-0016"}


def test_legacy_layout_with_data_below_the_skeleton_is_not_reshaped() -> None:
    """存量布局（骨架 13~24、数据 25+）：复用会把新行排到已有数据上方 ⇒ 不复用。"""
    physical = _fresh_physical()
    physical.update({25: "r1", 26: "r2"})
    candidates = {row: ident for row, ident in physical.items() if ident.startswith("GTROW")}
    adoption = S.plan_slot_adoption(
        candidates=candidates, physical=physical, wanted=["r1", "r2", "r3"], orphan=["r3"]
    )
    assert adoption is S.NO_ADOPTION


def test_skeleton_row_declared_by_the_projection_is_a_store_row_not_a_slot() -> None:
    physical = _fresh_physical()
    adoption = S.plan_slot_adoption(
        candidates=physical, physical=physical,
        wanted=["GTROW-D22-0013", "r1"], orphan=["r1"],
    )
    assert adoption.adopted == {"r1": 14}, "第 13 行是 store 已认领的骨架行，按身份照常更新"


def test_nothing_to_adopt_without_orphans_or_candidates() -> None:
    physical = _fresh_physical()
    assert S.plan_slot_adoption(
        candidates=physical, physical=physical, wanted=[], orphan=[]
    ) is S.NO_ADOPTION
    assert S.plan_slot_adoption(
        candidates={}, physical=physical, wanted=["r1"], orphan=["r1"]
    ) is S.NO_ADOPTION


# ═══════════════════════════════════════════════════════════════════════════
# 2. 候选（模板原样判据）与启用声明
# ═══════════════════════════════════════════════════════════════════════════


def _fv(row_key: str, column_key: str, value: Any, vt: ValueType, mode=FieldMode.editable):
    key = f"{_T}/{row_key}/{column_key}"
    return key, FieldValue(stable_key=key, value=value, value_type=vt, mode=mode, row_key=row_key)


def _scan(rows: dict[int, str]) -> Any:
    by_identity: dict[str, tuple[int, ...]] = {}
    for row, identity in sorted(rows.items()):
        by_identity[identity] = by_identity.get(identity, ()) + (row,)
    return SimpleNamespace(row_identity_by_row=rows, rows_by_identity=by_identity)


def _proj(*pairs: tuple[str, FieldValue]) -> Projection:
    return Projection(
        contract_id=P.PILOT_ADAPTER_ID, semantic_version="1.0.0", document_type="xlsx",
        values=dict(pairs),
    )


def test_d2_declares_reuse_for_the_detail_table_only() -> None:
    tables, baseline_fn = S.reuse_declaration(P.PILOT_ADAPTER_ID)
    assert tables == frozenset({P.ROWS_TABLE_KEY})
    assert baseline_fn is P.template_skeleton_baseline


def test_contracts_without_a_declaration_never_reuse() -> None:
    assert S.reuse_declaration("b60.hour_budget") == (frozenset(), None), "台账里有、但未声明"
    assert S.reuse_declaration("not.a.contract") == (frozenset(), None)
    assert S.skeleton_slot_candidates(
        contract_id="not.a.contract", table_key=_T, scan=_scan({13: _SK[0]}),
        substrate_projection=_proj(),
    ) == {}


def test_declaration_without_baseline_fails_closed(monkeypatch: pytest.MonkeyPatch) -> None:
    monkeypatch.setattr(P, "template_skeleton_baseline", None)
    with pytest.raises(S.SkeletonSlotReuseDeclarationError, match="template_skeleton_baseline"):
        S.reuse_declaration(P.PILOT_ADAPTER_ID)


def test_candidates_keep_template_rows_and_drop_user_filled_or_duplicated_rows() -> None:
    scan = _scan({13: _SK[0], 14: _SK[1], 15: _SK[2], 16: _SK[3], 17: _SK[3], 18: "r-real"})
    projection = _proj(
        _fv(_SK[0], "seq", 1, ValueType.integer),
        _fv(_SK[0], "credit_risk_classification", "单项计提", ValueType.enum),
        _fv(_SK[1], "customer_name", "用户在 OO 里填的客户", ValueType.text),
        # 受保护字段的缓存值不算用户内容
        _fv(_SK[2], "end_balance", 999, ValueType.amount, mode=FieldMode.formula),
    )
    candidates = S.skeleton_slot_candidates(
        contract_id=P.PILOT_ADAPTER_ID, table_key=_T, scan=scan, substrate_projection=projection
    )
    assert candidates == {13: _SK[0], 15: _SK[2]}, (
        "14 带用户内容、16/17 是重复身份、18 不是骨架 —— 都不得复用"
    )


def test_candidates_do_not_read_the_template_when_there_is_no_skeleton_row(
    monkeypatch: pytest.MonkeyPatch,
) -> None:
    def _must_not_be_called() -> SkeletonBaseline:
        raise AssertionError("没有骨架行时不得读模板")

    monkeypatch.setattr(P, "template_skeleton_baseline", _must_not_be_called)
    assert S.skeleton_slot_candidates(
        contract_id=P.PILOT_ADAPTER_ID, table_key=_T, scan=_scan({13: "r1"}),
        substrate_projection=_proj(),
    ) == {}


# ═══════════════════════════════════════════════════════════════════════════
# 3. 身份保留门：只豁免「声明退役 ∧ 骨架身份形态」
# ═══════════════════════════════════════════════════════════════════════════


def _inventories(expected_ids: list[str], observed_ids: list[str]) -> tuple[Any, Any]:
    expected = SimpleNamespace(
        defined_names=("GT_X",), table_present=True, table_ref="A12:AN24",
        row_uuids={str(13 + i): v for i, v in enumerate(expected_ids)},
        hidden_sheet_present=True,
    )
    observed = SimpleNamespace(
        defined_names=("GT_X",), table_present=True, table_ref="A12:AN30",
        identities=frozenset(observed_ids), hidden_sheet_present=True,
    )
    return expected, observed


def test_retention_exempts_only_the_declared_retired_skeletons() -> None:
    expected, observed = _inventories(_SK[:3], ["r1", "r2", _SK[2]])
    X.assert_identity_inventory_retained(
        expected=expected, observed=observed, entry_id="e",
        retired_row_identities=(_SK[0], _SK[1]),
    )
    with pytest.raises(X.IdentityRetentionError, match="丢失 1 个"):
        X.assert_identity_inventory_retained(
            expected=expected, observed=observed, entry_id="e",
            retired_row_identities=(_SK[0],),
        )


def test_retention_never_exempts_a_business_identity() -> None:
    expected, observed = _inventories(["dr-a", "dr-b"], ["dr-a"])
    with pytest.raises(X.IdentityRetentionError, match="'dr-b'"):
        X.assert_identity_inventory_retained(
            expected=expected, observed=observed, entry_id="e",
            retired_row_identities=("dr-b",),
        )


# ═══════════════════════════════════════════════════════════════════════════
# 4. 真链路：新 instrument 的 D2 权威模板
# ═══════════════════════════════════════════════════════════════════════════


def _d(label: str) -> str:
    return hashlib.sha256(label.encode("utf-8")).hexdigest()


@pytest.fixture(scope="module")
def contract() -> Any:
    return P.load_pilot_contract()


@pytest.fixture(scope="module")
def staged(contract: Any, tmp_path_factory: pytest.TempPathFactory) -> Any:
    return F.stage_instrumented_substrate(
        entry_id=P.PILOT_ENTRY_ID,
        staging_dir=tmp_path_factory.mktemp("d2-slot-reuse"),
        contract=contract,
    )


@pytest.fixture(scope="module")
def adapter(contract: Any, staged: Any) -> Any:
    def slot(kind: BundleSlot, digest: str) -> BundleSlotSpec:
        return BundleSlotSpec(kind, "definition", f"definition:{uuid.uuid4()}", digest)

    bundle = DefinitionBundleSnapshot(
        bundle_id=uuid.uuid4(),
        bundle_sha256=_d("d2-slot-reuse-bundle"),
        schema_version="definition-bundle:v1",
        state=DefinitionState.approved,
        authority_model=AuthorityModel.projection_contract,
        authority_model_definition_id=uuid.uuid4(),
        authority_model_definition_sha256=_d("d2-slot-reuse-authority"),
        slots={
            BundleSlot.template: slot(BundleSlot.template, contract.template_definition_sha256),
            BundleSlot.instrumentation: slot(
                BundleSlot.instrumentation, contract.instrumentation_definition_sha256
            ),
            BundleSlot.contract: slot(BundleSlot.contract, contract.canonical_sha256),
        },
    )
    definitions = FrozenEntryDefinitions(
        entry_id=P.PILOT_ENTRY_ID,
        bundle=bundle,
        contract=contract,
        adapter_build=AdapterBuild(
            adapter_id=P.PILOT_ADAPTER_ID,
            adapter_build_digest=_d("d2-slot-reuse-build"),
            document_type="xlsx",
            contract_version=str(contract.semantic_version),
        ),
        identity_inventory=staged.identity_inventory,
        business_sheets=(),
        dynamic_column_keys={},
        structure_inventory_size=0,
    )
    primary = F._identity_binding(provider=P, staged=staged, contract=contract)
    siblings = F._sibling_identity_bindings(
        provider=P, staged=staged, contract=contract, primary=primary
    )
    return AX.build_excel_adapter(
        definitions=definitions,
        binding=primary,
        direction="html_to_oo",
        sibling_bindings=siblings,
    )


def _store_rows(count: int) -> list[dict[str, Any]]:
    rows = []
    for index in range(count):
        rows.append({
            P.ROW_IDENTITY_STORE_KEY: f"dr-slot-{index:04d}",
            "seq": index + 1,
            "customerName": f"客户{index + 1:03d}",
            "companyCode": f"C{index + 1:03d}",
            "priorUnadjusted": float(1000 + index),
            "debitOccurrence": float(10 + index),
            "creditOccurrence": 1.0,
        })
    return rows


def _roundtrip(adapter: Any, contract: Any, staged: Any, out: Path, count: int) -> tuple[Any, Any]:
    store_projection = F._store_projection_for_provider(
        provider=P,
        store_payload={P.STORE_ITEM_ID: json.dumps(_store_rows(count), ensure_ascii=False)},
        contract=contract,
    )
    projection = F._overlay_store_on_substrate_baseline(
        adapter=adapter, contract=contract, substrate=staged.staged_path,
        store_projection=store_projection,
    )
    result = adapter.materialize(
        substrate=staged.staged_path, projection=projection, output=out, contract=contract
    )
    extracted = adapter.extract(
        artifact=out, contract=contract, retired_row_identities=result.retired_row_identities
    )
    assert_roundtrip_equivalent(intended=projection, extracted=extracted, contract=contract)
    adapter.verify_unmanaged_regions(
        before=staged.staged_path,
        after=out,
        contract=contract,
        row_shift=result.row_shift,
        total_formula_rows=result.total_formula_rows,
        propagation=result.workbook_row_change,
        per_table_shift=result.per_table_shift,
    ).assert_equivalent()
    return result, projection


def _sheet_column(out: Path, column: str, rows: range) -> list[Any]:
    import openpyxl

    wb = openpyxl.load_workbook(io.BytesIO(out.read_bytes()), read_only=True, data_only=False)
    try:
        ws = wb[P.MANAGED_SHEET]
        return [ws[f"{column}{row}"].value for row in rows]
    finally:
        wb.close()


def test_three_store_rows_start_at_row_13(adapter: Any, contract: Any, staged: Any, tmp_path: Path) -> None:
    out = tmp_path / "three.xlsx"
    result, _projection = _roundtrip(adapter, contract, staged, out, 3)
    assert result.retired_row_identities == tuple(_SK[:3])
    assert result.row_shift is None, "3 行全部落进骨架槽位，不需要插行"
    assert _sheet_column(out, P.UUID_COL, range(13, 17)) == [
        "dr-slot-0000", "dr-slot-0001", "dr-slot-0002", "GTROW-D22-0016",
    ]
    assert _sheet_column(out, "B", range(13, 16)) == ["客户001", "客户002", "客户003"]
    # 模板示例文字（AI13「单项计提」/ AI14~15「客户类型组合」）不得残留在已归属的业务行里
    assert _sheet_column(out, "AI", range(13, 16)) == [None, None, None]


def _cached_values(out: Path, column: str, rows: range) -> list[Any]:
    import openpyxl

    wb = openpyxl.load_workbook(io.BytesIO(out.read_bytes()), read_only=True, data_only=True)
    try:
        ws = wb[P.MANAGED_SHEET]
        return [ws[f"{column}{row}"].value for row in rows]
    finally:
        wb.close()


def test_adopted_rows_keep_formulas_but_drop_the_template_cached_values(
    adapter: Any, contract: Any, staged: Any, tmp_path: Path
) -> None:
    """复用行的公式格保留公式、去掉模板缓存值，反读与新插入行一致（store 没给值 ⇒ 无此键）。

    不去掉时模板的 `<v>0</v>` 反读成该行的值：首版记下的 projection 没这个键、发布产物有，
    同 store 再开在线编辑判「内容变了」，store 指纹直通永不命中（`test_store_materialize_pg`）。
    """
    out = tmp_path / "formulas.xlsx"
    result, projection = _roundtrip(adapter, contract, staged, out, 3)
    assert result.retired_row_identities == tuple(_SK[:3])
    # Q 期末余额 / S 本期未审数 / AB 本期审定数：三个公式列，本用例的 store 行都没给值
    for column in ("Q", "S", "AB"):
        assert all(str(f).startswith("=") for f in _sheet_column(out, column, range(13, 16))), column
        assert _cached_values(out, column, range(13, 16)) == [None, None, None], column
    assert _cached_values(out, "Q", range(16, 17)) == [0], "未被复用的骨架行不动"
    extracted = adapter.extract(artifact=out, contract=contract, retired_row_identities=result.retired_row_identities)
    for index in range(3):
        rid = f"dr-slot-{index:04d}"
        for leaf in ("end_balance", "current_unadjusted", "current_audited"):
            key = f"{_T}/{rid}/{leaf}"
            assert projection.get(key) is None and extracted.get(key) is None, key


def test_twenty_store_rows_fill_the_skeleton_then_insert_the_rest(
    adapter: Any, contract: Any, staged: Any, tmp_path: Path
) -> None:
    out = tmp_path / "twenty.xlsx"
    result, _projection = _roundtrip(adapter, contract, staged, out, 20)
    assert result.retired_row_identities == tuple(_SK)
    assert result.row_shift is not None and result.row_shift.count == 8
    assert result.row_shift.insert_at == 25
    assert _sheet_column(out, P.UUID_COL, range(13, 33)) == [
        f"dr-slot-{i:04d}" for i in range(20)
    ]
    assert _sheet_column(out, "A", range(33, 35)) == ["……", P.FOOTER_MARKER]


def test_undeclared_table_plan_is_identical_to_no_hook(
    adapter: Any, contract: Any, staged: Any, monkeypatch: pytest.MonkeyPatch
) -> None:
    """未声明启用时，挂上钩子与不挂钩子算出的计划逐项相同（逐字节零回归的计划侧判据）。"""
    monkeypatch.setattr(P, "SKELETON_SLOT_REUSE_TABLE_KEYS", ())
    S.clear_reuse_declaration_cache()
    store_projection = F._store_projection_for_provider(
        provider=P,
        store_payload={P.STORE_ITEM_ID: json.dumps(_store_rows(5), ensure_ascii=False)},
        contract=contract,
    )
    projection = F._overlay_store_on_substrate_baseline(
        adapter=adapter, contract=contract, substrate=staged.staged_path,
        store_projection=store_projection,
    )
    definitions = adapter.definitions
    binding = adapter.binding
    source = staged.staged_path.read_bytes()

    def plan(with_hook: bool) -> Any:
        step = M._plan_materialize_step(
            substrate=staged.staged_path, projection=projection, definitions=definitions,
            binding=binding, substrate_role=SubstrateRole.published_representation,
            substrate_kind=ArtifactKind.canonical, substrate_state=ArtifactState.published,
            source_bytes=source, capability=None, intended_formulas=None,
            limits=X.load_limits(), retain_identity_inventory=True,
        )
        if with_hook:
            return step.plan
        view = step.substrate_view
        import zipfile
        with zipfile.ZipFile(io.BytesIO(source)) as zf:
            runtime = X.read_runtime_binding_pairs(zf, metadata_sheet=binding.metadata_sheet)
        return M.plan_managed_writes(
            projection=projection, contract=definitions.contract, binding=binding,
            region=view.region, scan=view.scan, substrate_entries=M._read_entries(source),
            substrate_formulas=view.formula_inventory, runtime_binding=runtime,
        )

    hooked, bare = plan(True), plan(False)
    assert hooked.retired_row_identities == ()
    assert hooked.as_dict() == bare.as_dict()
    assert hooked.writes == bare.writes
    assert hooked.row_shift is not None and hooked.row_shift.insert_at == 25, (
        "未启用时保持旧行为：追加在骨架之后"
    )


def test_verify_before_digest_cache_is_keyed_by_the_managed_region(
    adapter: Any, contract: Any, staged: Any, tmp_path: Path, monkeypatch: pytest.MonkeyPatch
) -> None:
    """同一个 before 先后对 3 行 / 20 行插入做未管理区域校验，两次都必须等价。

    受管坐标由 region（从 **after** 解析）决定；before digest 缓存键缺 region 时，第二次会拿到
    第一次的 digest ⇒ 假漂移。与复用无关（关掉复用、纯插行同样复现），故单独钉住。
    """
    from app.services.workpaper_sync.parse_cache import BEFORE_DIGEST_CACHE

    monkeypatch.setattr(P, "SKELETON_SLOT_REUSE_TABLE_KEYS", ())
    S.clear_reuse_declaration_cache()
    BEFORE_DIGEST_CACHE.clear()
    for count in (3, 20):
        result, _projection = _roundtrip(adapter, contract, staged, tmp_path / f"seq{count}.xlsx", count)
        assert result.retired_row_identities == ()
        assert result.row_shift is not None and result.row_shift.count == count

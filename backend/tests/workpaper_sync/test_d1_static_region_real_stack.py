# -*- coding: utf-8 -*-
"""D1-4 第三区静态受管区的**真栈**写入 —— 真模板 + 真 instrumentation + 真 plan/apply。

spec: d1-sync-row-table-engine-and-d1-coverage · 契约层 `static_tables` 收口

═══ 与 `test_d1_static_region_contract.py` 的分工 ═══

那份是**纯函数层**（契约形态 / 投影 / 回写逐值），不碰 xlsx 字节。
本份补的是它明确登记为「仍未闭合」的那条：**真 xlsx 上的写入行为**
——`plan_managed_writes` 的 static_tables 分支、`apply_plan_zip`、
以及 `verify_unmanaged_regions` 对「未声明的列前后相同」的判定。

═══ 走哪条路径：生产走的是**同 sheet 动态 binding** ═══

`plan_managed_writes` 的分派是 `if is_static_region(binding): return _plan_static_writes(...)`，
否则走动态路径 —— 而动态路径**自己也遍历 static_tables**（`excel_materialize.py:1942`）。

D1 的生产 binding 集合里**没有**静态 binding：`_static_region_bindings` 要求静态声明里带
`tables[0].table_key`，而 `phase5_d1_expansion._static_sheet_declarations()` 不带
（同 sheet 动静并列不需要独立 binding）⇒ 第三区的格由 `d14-managed` 的**动态** binding 写。

⇒ 本文件刻意用动态 binding 跑，不用静态 binding —— 测生产不走的那条路等于假绿。

🔴 由此引出一个必须实测的问题：`d14-managed` 上有**两个**动态 binding
（individual / portfolio），两者的 `managed_tables_of` 都会返回同一张 static table
⇒ 第三区的格会被**写两次**。见 `test_two_dynamic_bindings_write_static_cells_idempotently`。
"""

from __future__ import annotations

import io
import json
import sys
import zipfile
from functools import lru_cache
from pathlib import Path

import openpyxl
import pytest

_BACKEND = Path(__file__).resolve().parents[2]
if str(_BACKEND) not in sys.path:
    sys.path.insert(0, str(_BACKEND))

from app.services.workpaper_sync import excel_instrumentation as EI  # noqa: E402
from app.services.workpaper_sync import excel_materialize as M  # noqa: E402
from app.services.workpaper_sync import phase5_d1_04_bad_debt as D104  # noqa: E402
from app.services.workpaper_sync import phase5_d1_notes_receivable as P  # noqa: E402
from app.services.workpaper_sync.contracts import parse_contract  # noqa: E402
from app.services.workpaper_sync.excel_extract import (  # noqa: E402
    ExcelIdentityBinding,
    is_static_region,
)

from app.services.workpaper_sync import phase5_d1_expansion as EXP  # noqa: E402

# 🔴 T7 裁决 A：静态 table 已撤回（`_INCLUDE_D104_NOTETYPE_STATIC=False`）。
#    本文件是「静态 table **在**契约里」时的真栈判据 ⇒ 通路不存在时整体跳过，**不删** ——
#    marker-relative content 字段落地、开关翻回 True 时自动复活。
#    排除态的正向守卫在 `test_d104_static_region_excluded.py`（含失效条目反向检查）。
pytestmark = pytest.mark.skipif(
    not EXP._INCLUDE_D104_NOTETYPE_STATIC,
    reason="T7 裁决 A：D1-4 静态 table 已撤回；排除态守卫见 test_d104_static_region_excluded.py",
)

SHEET = D104.MANAGED_SHEET_D104
TABLE_KEY = D104.SPEC_D104_NOTETYPE.table_key
ITEM_ID = D104.SPEC_D104_NOTETYPE.store_item_id

#: 投影会写的 6 列 → 期望坐标（2 固定行）
PROJECTED_COORDS = {
    "A23", "B23", "C23", "D23", "L23", "M23",
    "A24", "B24", "C24", "D24", "L24", "M24",
}
#: 模板公式列（投影不供 ⇒ 不该出现写入）
FORMULA_COORDS = {"E23", "K23", "N23", "E24", "K24", "N24"}

#: HTML 不拥有的列（模板 F~J「计提/其他增加/转回/核销/其他减少」）。
#:
#: 🔴🔴 **刻意写成字面量，不读生产常量 `_HTML_UNOWNED_COLUMNS_D104_NOTETYPE`**。
#:    本文件第一版就是遍历那个常量的 —— 变异实测（把 `F` 从常量里拿掉）时
#:    **判据跟着不检查 F 了、11 条全绿**，正是本仓库反复出现的
#:    「测试镜像同款错误 ⇒ 恒绿而生产恒死」。
#:    守卫的期望值必须独立于被守对象；另有一条判据专门断言
#:    「生产常量 == 本字面量」，常量改了会在那里打红而不是在这里静默失效。
HTML_UNOWNED_COLUMNS = ("F", "G", "H", "I", "J")


@lru_cache(maxsize=1)
def _instrumented() -> bytes:
    """真 D1 模板 + 真 instrumentation（秒级，进程内只做一次）。"""
    # 🔴 `instrumentation_specs()`（复数）在**伴生模块** `phase5_d1_expansion` 上，
    #    entry 模块只有单数 `instrumentation_spec()`（D1-3 那一张）。
    #    按 entry 模块取会被 PEP 562 `__getattr__` 挡下并提示「Did you mean: instrumentation_spec」。
    from app.services.workpaper_sync import phase5_d1_expansion as _exp

    return EI.instrument_workbook_bytes_multi(
        P.read_authoritative_template(),
        _exp.instrumentation_specs(),
        gate=P.excel_carrier_gate(),
    ).instrumented_bytes


@pytest.fixture(scope="module")
def contract():
    return parse_contract(P.build_contract_payload(), adapter_id=P.ADAPTER_ID)


@pytest.fixture(scope="module")
def substrate() -> bytes:
    return _instrumented()


def _payload() -> str:
    rows = []
    for i, (row_id, _row_no, label) in enumerate(D104.NOTETYPE_FIXED_ROWS, start=1):
        rows.append({
            "rowId": row_id, "noteType": label, "isFixed": True,
            "priorUnadjusted": 1000.0 * i, "priorAje": 11.0 * i, "priorRje": 22.0 * i,
            "currentUnadjusted": 2000.0 * i, "currentAje": 33.0 * i, "currentRje": 44.0 * i,
        })
    return json.dumps(rows, ensure_ascii=False)


def _dynamic_binding(spec) -> ExcelIdentityBinding:
    return ExcelIdentityBinding(
        table_key=spec.table_key,
        table_name=spec.table_name,
        uuid_column=spec.uuid_col,
        metadata_sheet="_GT_SYNC",
    )


def _real_scan(spec, *, contract, substrate: bytes):
    """从**真 instrumented 工作簿**读该动态区的隐藏 UUID 列，造真 `RowIdentityScan`。

    🔴 不能传 `scan=None` —— 那是**静态** binding 才允许的（`_plan_static_writes` 用不到它）；
       动态路径会 `scan.row_identity_by_row` ⇒ `AttributeError: 'NoneType'`。
       本文件走的是动态 binding（生产就是这条），所以必须给真 scan。
    """
    import hashlib

    from app.services.workpaper_sync.excel_extract import _scan_row_identities

    ws = openpyxl.load_workbook(io.BytesIO(substrate), data_only=False)[SHEET]
    raw_by_row = {
        row: str(ws[f"{spec.uuid_col}{row}"].value or "")
        for row in range(spec.first_data_row, spec.last_data_row + 1)
    }
    table = next(
        t for s in contract.sheets for t in s.tables if t.table_key == spec.table_key
    )
    return _scan_row_identities(
        table=table,
        sheet_name=SHEET,
        uuid_column=spec.uuid_col,
        raw_by_row=raw_by_row,
        artifact_sha256=hashlib.sha256(substrate).hexdigest(),
        tombstoned=(),
    )


def _runtime_binding(substrate: bytes):
    """从真 instrumented 工作簿的 `_GT_SYNC` 隐藏 sheet 读冻结预期。"""
    from app.services.workpaper_sync.excel_extract import read_runtime_binding_pairs

    with zipfile.ZipFile(io.BytesIO(substrate)) as zf:
        return read_runtime_binding_pairs(zf)


def _plan_for(spec, *, contract, substrate: bytes):
    proj = D104.build_notetype_store_projection(_payload(), contract=contract)
    binding = _dynamic_binding(spec)
    assert not is_static_region(binding), "本文件刻意用**动态** binding（生产就是这条）"
    with zipfile.ZipFile(io.BytesIO(substrate)) as zf:
        region = M.resolve_managed_region(zf, contract=contract, binding=binding)
    return M.plan_managed_writes(
        projection=proj,
        contract=contract,
        binding=binding,
        region=region,
        scan=_real_scan(spec, contract=contract, substrate=substrate),
        substrate_entries=M._read_entries(substrate),
        substrate_formulas={},
        # 🔴 动态路径有 footer 两门，要 `GT_FOOTER_ROW` 的冻结预期 —— 传 `{}` 会得
        #    `FooterAnchorDriftError: runtime binding 里没有 GT_FOOTER_ROW`。
        #    读侧**唯一**入口是 `read_runtime_binding_pairs`（不得在别处抄第二份 XML 解析）。
        runtime_binding=_runtime_binding(substrate),
    ), binding


# ═══════════════════════════════════════════════════════════════════════════
# 1. 真模板 instrumentation：definedName 锚点真的注进去了
# ═══════════════════════════════════════════════════════════════════════════


def test_instrumentation_injects_static_defined_name(substrate: bytes) -> None:
    """第三区的 workbook-scope definedName 必须在真模板上注入成功。

    注不进去 ⇒ 后续 `resolve_managed_region` 找不到锚点，整条静态通路在真栈上不可达
    （纯函数层的判据完全看不到这一步）。
    """
    with zipfile.ZipFile(io.BytesIO(substrate)) as zf:
        wb_xml = zf.read("xl/workbook.xml").decode("utf-8")
    name = D104.SPEC_D104_NOTETYPE.defined_name
    assert f'name="{name}"' in wb_xml, f"definedName {name} 未注入真模板"


def test_template_formulas_exist_before_write(substrate: bytes) -> None:
    """E/K/N 三列在真模板上确有公式 —— 这是把它们声明为 `formula` 的前提。

    契约声明 formula 而 substrate 没公式时 `_emit` 会抛 `ProtectedRegionWriteError`，
    本条先把前提钉住，否则后面的「公式保留」断言无从谈起。
    """
    wb = openpyxl.load_workbook(io.BytesIO(substrate), data_only=False)
    ws = wb[SHEET]
    for coord in sorted(FORMULA_COORDS):
        val = ws[coord].value
        assert isinstance(val, str) and val.startswith("="), (
            f"{coord} 在真模板上不是公式（实得 {val!r}）"
        )


# ═══════════════════════════════════════════════════════════════════════════
# 2. plan：写哪些格、不写哪些格
# ═══════════════════════════════════════════════════════════════════════════


def test_plan_writes_exactly_the_projected_static_cells(contract, substrate) -> None:
    """🔴 真栈写入面 = 投影的 6 列 × 2 固定行，一格不多一格不少。"""
    plan, _b = _plan_for(D104.SPEC_D104_INDIVIDUAL, contract=contract, substrate=substrate)
    coords = {w.coord for w in plan.writes}
    assert PROJECTED_COORDS <= coords, f"缺写：{sorted(PROJECTED_COORDS - coords)}"
    # 第三区之外的格不该被本次投影碰到（投影里只有第三区的键）
    third_region = {c for c in coords if c[1:] in ("23", "24")}
    assert third_region == PROJECTED_COORDS, (
        f"第三区写入面不符，多写={sorted(third_region - PROJECTED_COORDS)}"
    )


def test_plan_does_not_touch_formula_cells(contract, substrate) -> None:
    """🔴 E/K/N 不在投影里 ⇒ `_emit` 跳过 ⇒ plan 里不该有它们的写入。

    这条是「投影只供 HTML 拥有的列」在真栈上的落点：若哪天把公式列塞进投影，
    这里会立刻打红（纯函数层只能看到投影，看不到 plan）。
    """
    plan, _b = _plan_for(D104.SPEC_D104_INDIVIDUAL, contract=contract, substrate=substrate)
    coords = {w.coord for w in plan.writes}
    assert not (FORMULA_COORDS & coords), (
        f"公式格被写了：{sorted(FORMULA_COORDS & coords)}"
    )


def test_plan_has_no_row_shift_or_footer_gate_for_static_region(contract, substrate) -> None:
    """静态区绕开位移链：本次投影只含第三区 ⇒ 不产生 row_shift。"""
    plan, _b = _plan_for(D104.SPEC_D104_INDIVIDUAL, contract=contract, substrate=substrate)
    assert plan.row_shift is None, "静态区不该触发行位移"


# ═══════════════════════════════════════════════════════════════════════════
# 3. apply → 真字节上复读
# ═══════════════════════════════════════════════════════════════════════════


@pytest.fixture(scope="module")
def applied(contract, substrate) -> bytes:
    plan, _b = _plan_for(D104.SPEC_D104_INDIVIDUAL, contract=contract, substrate=substrate)
    return M.apply_plan_zip(substrate, plan)


def test_values_land_in_the_right_cells(applied: bytes) -> None:
    """🔴 真字节复读：值落在正确坐标，两行不串。"""
    wb = openpyxl.load_workbook(io.BytesIO(applied), data_only=False)
    ws = wb[SHEET]
    assert ws["A23"].value == "银行承兑汇票小计"
    assert ws["A24"].value == "商业承兑汇票小计"
    assert float(ws["B23"].value) == 1000.0
    assert float(ws["B24"].value) == 2000.0
    assert float(ws["C23"].value) == 11.0 and float(ws["C24"].value) == 22.0
    assert float(ws["L23"].value) == 33.0 and float(ws["L24"].value) == 66.0
    assert float(ws["M23"].value) == 44.0 and float(ws["M24"].value) == 88.0


def test_template_formulas_survive_the_write(substrate: bytes, applied: bytes) -> None:
    """🔴 E/K/N 的公式**逐字未变** —— 这是「分权拥有」在真栈上的证明。

    F~J 由审计师在 Excel 填、K 由模板公式从它们算 ⇒ 我们把公式写坏了就等于毁掉这条语义。
    """
    before = openpyxl.load_workbook(io.BytesIO(substrate), data_only=False)[SHEET]
    after = openpyxl.load_workbook(io.BytesIO(applied), data_only=False)[SHEET]
    for coord in sorted(FORMULA_COORDS):
        assert after[coord].value == before[coord].value, (
            f"{coord} 的公式被改写了：{before[coord].value!r} → {after[coord].value!r}"
        )


def test_html_unowned_columns_are_untouched(substrate: bytes, applied: bytes) -> None:
    """🔴 F~J（HTML 不拥有）前后**逐格相同** —— 若哪天把它们声明回去，
    `_render_number(None)` 会把它们写成 `0`，这条会立刻打红。"""
    before = openpyxl.load_workbook(io.BytesIO(substrate), data_only=False)[SHEET]
    after = openpyxl.load_workbook(io.BytesIO(applied), data_only=False)[SHEET]
    for col in HTML_UNOWNED_COLUMNS:  # 字面量，不读生产常量（见其注释）
        for row in (23, 24):
            coord = f"{col}{row}"
            assert after[coord].value == before[coord].value, (
                f"{coord} 被改动了（HTML 不拥有该列，应原样保留）"
            )


def test_html_unowned_columns_are_not_declared_in_the_contract(contract) -> None:
    """🔴 F~J 一列都不许出现在契约的静态 table 里。

    这条与上一条是**两件事**：上一条验「没被写」，本条验「没被声明」。
    只验前者不够 —— 声明了但当前投影不供，反向（extract/merge）仍会把这些 Excel 格
    读回来并塞进 HTML 行对象里前端从不读的键；而一旦哪天有人把它们加进投影，
    `_render_number(None)` 就会把 F23:J24 写成 `0`。

    🔴 期望值用**字面量** `HTML_UNOWNED_COLUMNS`，不读生产常量 —— 读常量会随变异一起漂。
    """
    t = next(t for s in contract.sheets for t in s.tables if t.table_key == TABLE_KEY)
    declared = {f.cell.column for f in t.fields if f.cell}
    offenders = sorted(set(HTML_UNOWNED_COLUMNS) & declared)
    assert not offenders, (
        f"HTML 不拥有的列被声明进契约：{offenders} —— materialize 会把它们写成 0/清空"
    )


def test_production_constant_matches_the_literal_expectation() -> None:
    """生产常量 ↔ 本文件字面量**逐值相等**。

    这是把「期望值独立于被守对象」这件事闭上环：常量若被改动，红在**这里**
    （一眼看出是常量变了），而不是让上面两条判据静默跟着漂。
    """
    assert D104._HTML_UNOWNED_COLUMNS_D104_NOTETYPE == frozenset(HTML_UNOWNED_COLUMNS), (
        f"生产常量 {sorted(D104._HTML_UNOWNED_COLUMNS_D104_NOTETYPE)} 与本文件字面量 "
        f"{sorted(HTML_UNOWNED_COLUMNS)} 不一致 —— 改动了就要在这里对账，不能让守卫跟着漂"
    )


# ═══════════════════════════════════════════════════════════════════════════
# 4. verify_unmanaged_regions：未声明的格前后相同 ⇒ 判等价
# ═══════════════════════════════════════════════════════════════════════════


def test_unmanaged_regions_stay_equivalent(contract, substrate, applied, tmp_path) -> None:
    """🔴 真栈校验门：只动受管格 ⇒ `verify_unmanaged_regions` 必须判 equivalent。

    这条同时证明「F~J 不进 field_specs」不会被校验门反过来判成漂移
    （纯函数层推理过，但没在真字节上验过）。
    """
    from app.services.workpaper_sync.excel_extract import verify_unmanaged_regions

    before_p = tmp_path / "before.xlsx"
    after_p = tmp_path / "after.xlsx"
    before_p.write_bytes(substrate)
    after_p.write_bytes(applied)

    binding = _dynamic_binding(D104.SPEC_D104_INDIVIDUAL)
    with zipfile.ZipFile(io.BytesIO(substrate)) as zf:
        region = M.resolve_managed_region(zf, contract=contract, binding=binding)
    report = verify_unmanaged_regions(
        before=before_p, after=after_p, contract=contract, region=region, binding=binding
    )
    assert report.equivalent, f"未受管区被判漂移：{getattr(report, 'first_difference', None)}"


# ═══════════════════════════════════════════════════════════════════════════
# 5. 🔴 同 sheet 两个动态 binding 都会写第三区 —— 幂等性必须实测
# ═══════════════════════════════════════════════════════════════════════════


def test_two_dynamic_bindings_write_static_cells_idempotently(contract, substrate) -> None:
    """🔴 `d14-managed` 有 individual / portfolio 两个动态 binding，
    两者的 `managed_tables_of` 都返回同一张 static table ⇒ 第三区的格**会被写两次**。

    这不是缺陷但必须钉住**幂等**：两次写入的坐标与载荷逐项相同 ⇒ 先后顺序无关、
    也不会触发跨 binding 载荷冲突。若哪天两条路径算出不同值（例如一方漏了切片），
    本条立刻打红 —— 这正是「只验接线不验语义会漏掉什么」的那类缺陷。
    """
    plan_a, _ = _plan_for(D104.SPEC_D104_INDIVIDUAL, contract=contract, substrate=substrate)
    plan_b, _ = _plan_for(D104.SPEC_D104_PORTFOLIO, contract=contract, substrate=substrate)

    def third_region(plan):
        return {
            w.coord: (w.kind, w.value)
            for w in plan.writes
            if w.coord[1:] in ("23", "24")
        }

    a, b = third_region(plan_a), third_region(plan_b)
    assert set(a) == PROJECTED_COORDS and set(b) == PROJECTED_COORDS, (
        f"两个 binding 的第三区写入面不等：only_a={sorted(set(a) - set(b))} "
        f"only_b={sorted(set(b) - set(a))}"
    )
    assert a == b, (
        "两个 binding 对第三区算出的载荷不同 ⇒ 写入顺序会影响结果："
        + str({k: (a[k], b[k]) for k in a if a[k] != b[k]})
    )


def test_applying_twice_is_byte_stable(contract, substrate) -> None:
    """再验一层：同一 plan 连应用两次，第三区的值不漂。"""
    plan, _b = _plan_for(D104.SPEC_D104_INDIVIDUAL, contract=contract, substrate=substrate)
    once = M.apply_plan_zip(substrate, plan)
    twice = M.apply_plan_zip(once, plan)
    ws1 = openpyxl.load_workbook(io.BytesIO(once), data_only=False)[SHEET]
    ws2 = openpyxl.load_workbook(io.BytesIO(twice), data_only=False)[SHEET]
    for coord in sorted(PROJECTED_COORDS | FORMULA_COORDS):
        assert ws1[coord].value == ws2[coord].value, f"{coord} 二次应用后漂了"

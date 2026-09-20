# -*- coding: utf-8 -*-
"""D4-9 instrumentation 真实注入 + sibling binding 对齐 —— 行为级验证（judge-first）。

spec: d4-9-customer-structure-bidirectional-writeback · Task 5
Requirements 1.1/1.4/1.5

架构（路 B）：D4-9 两 instrumentation spec（W/X）由 sibling provider
phase5_d4_customer_structure.instrumentation_spec_d49 提供，父模块
phase5_d4_revenue_detail.instrumentation_specs() 编排。

判据（真实注入 / 真实 parse_contract，非符号存在）：
1. instrument_workbook_bytes_multi 对 D4-9 两 spec（W/X）→ openpyxl 同时认出
   GT_D49C_ROWS + GT_D49P_ROWS（同 sheet 双区）。
2. 父 instrumentation_specs() 含 D4-9 两 spec；契约 D4-9 sheet 恰 2 张行 table
   （totals row_identity=None 不计入），UUID 列 W/X 互不相同。
3. sibling binding 对齐：D4-9 两 spec 对齐到 current/prior 两行 table，不串区。
"""
from __future__ import annotations

import io

import pytest

from app.services.workpaper_sync import phase5_d4_customer_structure as m
from app.services.workpaper_sync import phase5_d4_revenue_detail as parent
from app.services.workpaper_sync.contracts import parse_contract
from app.services.workpaper_sync.excel_instrumentation import (
    instrument_workbook_bytes_multi,
)


def _tables_on_sheet(instrumented_bytes: bytes) -> list[str]:
    from openpyxl import load_workbook

    wb = load_workbook(io.BytesIO(instrumented_bytes))
    try:
        return sorted(wb[m.MANAGED_SHEET_D49].tables.keys())
    finally:
        wb.close()


def test_d49_two_specs_real_injection_yields_both_tables() -> None:
    """只注入 D4-9 两 spec（隔离验证同 sheet 双区注入内核对 D4-9 生效）。"""
    specs = m.instrumentation_spec_d49(
        entry_id=parent.ENTRY_ID, template_relative_path=parent.TEMPLATE_RELATIVE_PATH
    )
    inst = instrument_workbook_bytes_multi(
        parent.read_authoritative_template(), specs, gate=parent.excel_carrier_gate()
    )
    assert _tables_on_sheet(inst.instrumented_bytes) == [
        m.TABLE_NAME_CURRENT,
        m.TABLE_NAME_PRIOR,
    ]


def test_parent_instrumentation_specs_include_d49() -> None:
    specs = parent.instrumentation_specs()
    d49_specs = [s for s in specs if str(s.managed_sheet) == m.MANAGED_SHEET_D49]
    assert len(d49_specs) == 2, f"父 instrumentation_specs 应含 D4-9 两 spec，实得 {len(d49_specs)}"
    uuid_cols = sorted(str(s.uuid_col) for s in d49_specs)
    assert uuid_cols == ["W", "X"], f"D4-9 两 spec UUID 列应为 W/X，实得 {uuid_cols}"


def test_d49_sheet_has_two_row_tables_totals_excluded() -> None:
    contract = parse_contract(parent.build_contract_payload(), adapter_id=parent.ADAPTER_ID)
    d49 = {s.sheet_key: s for s in contract.sheets}[m.SHEET_KEY_D49]
    row_tables = [t.table_key for t in d49.tables if t.row_identity is not None]
    assert row_tables == [m.ROWS_TABLE_KEY_CURRENT, m.ROWS_TABLE_KEY_PRIOR], (
        f"D4-9 行 table（row_identity 非空）应恰 2 张，totals 不计入，实得 {row_tables}"
    )


def test_mapping_digest_frozen() -> None:
    assert m.assert_mapping_digest_d49() == m.EXPECTED_MAPPING_DIGEST_D49

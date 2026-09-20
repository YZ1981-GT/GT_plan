"""静态受管区（definedName-anchored static cell sheet）双向回写守卫。

spec: workpaper-sync-static-cell-sheet-writeback
覆盖 Property 1/4/5/6（静态往返等价 / 锚点缺失 fail-closed / kind 显式分派 / 无幽灵 UUID）。

本文件按 spec 实现顺序**逐 wave 追加**：Task 1 先钉 binding kind 判据 + 静态形态构造/校验，
后续 task 追加 resolve/extract/materialize 往返与 observer 分派。
"""

from __future__ import annotations

import pytest

from app.services.workpaper_sync import excel_extract as X
from app.services.workpaper_sync.excel_extract import (
    BindingKind,
    ExcelIdentityBinding,
    ManagedRegionResolutionError,
    is_static_region,
)


# ═══════════════════════════════════════════════════════════════════════════
# Task 1 · binding kind 判据 + 静态形态（Requirement 1 / Property 5）
# ═══════════════════════════════════════════════════════════════════════════


class TestBindingKind:
    """kind 由 `defined_name` 存在决定，不是 uuid_column 缺失（Property 5）。"""

    def test_dynamic_binding_kind_is_excel_table(self) -> None:
        binding = ExcelIdentityBinding(
            table_key="k_rows", table_name="GT_TBL", uuid_column="Z"
        )
        assert binding.kind is BindingKind.excel_table
        assert is_static_region(binding) is False

    def test_static_binding_kind_is_static_region(self) -> None:
        binding = ExcelIdentityBinding(
            table_key="d433-static", defined_name="GT_MANAGED_REGION_D433"
        )
        assert binding.kind is BindingKind.static_region
        assert is_static_region(binding) is True

    def test_static_binding_omits_table_name_and_uuid(self) -> None:
        # 静态构造可省略 table_name/uuid_column（默认 ""），不因缺它们而抛。
        binding = ExcelIdentityBinding(
            table_key="d433-static", defined_name="GT_MANAGED_REGION_D433"
        )
        assert binding.table_name == ""
        assert binding.uuid_column == ""
        assert binding.defined_name == "GT_MANAGED_REGION_D433"


class TestBindingKindFailClosed:
    """混填 / 空必填字段 fail-closed（Requirement 1.1/1.2）。"""

    def test_static_binding_rejects_table_name(self) -> None:
        with pytest.raises(ManagedRegionResolutionError, match="kind 混填"):
            ExcelIdentityBinding(
                table_key="d433-static",
                defined_name="GT_MANAGED_REGION_D433",
                table_name="GT_TBL",  # 混填
            )

    def test_static_binding_rejects_uuid_column(self) -> None:
        with pytest.raises(ManagedRegionResolutionError, match="kind 混填"):
            ExcelIdentityBinding(
                table_key="d433-static",
                defined_name="GT_MANAGED_REGION_D433",
                uuid_column="Z",  # 混填
            )

    def test_static_binding_rejects_empty_table_key(self) -> None:
        with pytest.raises(ManagedRegionResolutionError, match="table_key"):
            ExcelIdentityBinding(
                table_key="", defined_name="GT_MANAGED_REGION_D433"
            )

    def test_static_binding_rejects_empty_metadata_sheet(self) -> None:
        with pytest.raises(ManagedRegionResolutionError, match="metadata_sheet"):
            ExcelIdentityBinding(
                table_key="d433-static",
                defined_name="GT_MANAGED_REGION_D433",
                metadata_sheet="",
            )


class TestDynamicBindingUnchanged:
    """动态 binding 的既有强校验逐条保留（Requirement 1.5，动态零回归）。"""

    def test_dynamic_rejects_empty_table_name(self) -> None:
        with pytest.raises(ManagedRegionResolutionError, match="table_name"):
            ExcelIdentityBinding(table_key="k_rows", uuid_column="Z")

    def test_dynamic_rejects_empty_uuid_column(self) -> None:
        with pytest.raises(ManagedRegionResolutionError, match="uuid_column"):
            ExcelIdentityBinding(table_key="k_rows", table_name="GT_TBL")

    def test_dynamic_rejects_malformed_uuid_column(self) -> None:
        with pytest.raises(ManagedRegionResolutionError, match="形态非法"):
            ExcelIdentityBinding(
                table_key="k_rows", table_name="GT_TBL", uuid_column="123"
            )

    def test_dynamic_accepts_valid(self) -> None:
        binding = ExcelIdentityBinding(
            table_key="k_rows", table_name="GT_TBL", uuid_column="AA"
        )
        assert binding.kind is BindingKind.excel_table
        assert binding.uuid_column == "AA"


# ═══════════════════════════════════════════════════════════════════════════
# Task 2 · resolve_managed_region definedName 分支（Requirement 2 / Property 4）
# ═══════════════════════════════════════════════════════════════════════════

import io
import zipfile

import openpyxl
from openpyxl.workbook.defined_name import DefinedName

from app.services.workpaper_sync.contracts import parse_contract
from app.services.workpaper_sync.excel_extract import (
    IdentityCarrierMissingError,
    _resolve_static_region,
    _split_defined_name_ref,
)

STATIC_SHEET = "其他业务毛利率分析表D433"
STATIC_DEFINED_NAME = "GT_MANAGED_REGION_D433"
STATIC_REF = "$E$12:$M$13"  # 迷你几何（2 行 × 若干列），Task 2 只验区域求解
STATIC_TABLE_KEY = "d433-static"
STATIC_CONTRACT_ID = "d4.other_margin_static.test"


def _make_static_workbook(
    *, defined_name: str = STATIC_DEFINED_NAME, ref: str = STATIC_REF,
    sheet_title: str = STATIC_SHEET, workbook_scope: bool = True,
    duplicate: bool = False, on_hidden_meta: bool = False,
) -> bytes:
    """造一个含 workbook-scope definedName 的最小静态 sheet workbook。"""
    from app.services.excel_structure_fingerprint import GT_SYNC_SHEET_NAME

    wb = openpyxl.Workbook()
    ws = wb.active
    ws.title = sheet_title
    for r in range(12, 14):
        for c in ("E", "F", "H", "I", "K", "L"):
            ws[f"{c}{r}"] = 0
    target_title = sheet_title
    if on_hidden_meta:
        meta = wb.create_sheet(GT_SYNC_SHEET_NAME)
        meta["A1"] = "k"
        target_title = GT_SYNC_SHEET_NAME
    full_ref = f"'{target_title}'!{ref}"
    if workbook_scope:
        wb.defined_names[defined_name] = DefinedName(defined_name, attr_text=full_ref)
        if duplicate:
            # 再插一个同名 workbook-scope（openpyxl dict 会去重，故直接改 XML 注入第二个）
            pass
    else:
        # sheet-local scope（localSheetId）
        dn = DefinedName(defined_name, attr_text=full_ref, localSheetId=0)
        wb[sheet_title].defined_names[defined_name] = dn
    buf = io.BytesIO()
    wb.save(buf)
    data = buf.getvalue()
    if duplicate:
        data = _inject_duplicate_defined_name(data, defined_name, full_ref)
    return data


def _inject_duplicate_defined_name(data: bytes, name: str, ref: str) -> bytes:
    """往 workbook.xml 的 definedNames 里再塞一个同名 workbook-scope name。"""
    with zipfile.ZipFile(io.BytesIO(data)) as zf:
        parts = {n: zf.read(n) for n in zf.namelist()}
    wbxml = parts["xl/workbook.xml"].decode("utf-8")
    extra = f'<definedName name="{name}">{ref}</definedName>'
    if "<definedNames>" in wbxml:
        wbxml = wbxml.replace("<definedNames>", "<definedNames>" + extra, 1)
    else:
        wbxml = wbxml.replace("</sheets>", "</sheets><definedNames>" + extra + "</definedNames>", 1)
    parts["xl/workbook.xml"] = wbxml.encode("utf-8")
    out = io.BytesIO()
    with zipfile.ZipFile(out, "w", zipfile.ZIP_DEFLATED) as zf:
        for n, b in parts.items():
            zf.writestr(n, b)
    return out.getvalue()


def _static_contract():
    import hashlib

    def _d(s: str) -> str:
        return hashlib.sha256(s.encode()).hexdigest()

    payload = {
        "schema_version": "contract-definition:v1",
        "contract_id": STATIC_CONTRACT_ID,
        "semantic_version": "1.0.0",
        "review_status": "reviewed",
        "document_type": "xlsx",
        "template_definition_sha256": _d("d433-template-definition"),
        "instrumentation_definition_sha256": _d("d433-instrumentation-definition"),
        "template": {
            "relative_path": "D/其他业务毛利率分析表D433.xlsx",
            "template_sha256": _d("d433-template"),
            "normalized_structure_hash": _d("d433-normalized-structure"),
        },
        "identity_carriers": ["defined_name"],
        "sheets": [
            {
                "sheet_key": "d433-managed",
                "excel_name": STATIC_SHEET,
                # locator anchor 复用已探针通过的 defined_name_ref（真实 OO 已证 passed）；
                # 「静态 vs 转置」由 binding.kind + observer payload region_kind 区分，非 carrier 锚点名。
                "locator": {"anchor": "defined_name_ref"},
                "tables": [
                    {
                        "table_key": STATIC_TABLE_KEY,
                        "anchor": "E12",
                        "header_rows": 1,
                        "fields": [
                            {
                                "stable_field_key": f"{STATIC_TABLE_KEY}/rent_rev_m1",
                                "json_pointer": "/months/0/revenue",
                                "column_key": "rent_rev_m1",
                                "cell": {"column": "E", "row_from": 12},
                                "mode": "editable",
                                "value_type": "amount",
                                "source_ref": "wp:D4-33!E12",
                            },
                            {
                                "stable_field_key": f"{STATIC_TABLE_KEY}/rent_cost_m1",
                                "json_pointer": "/months/0/cost",
                                "column_key": "rent_cost_m1",
                                "cell": {"column": "F", "row_from": 12},
                                "mode": "editable",
                                "value_type": "amount",
                                "source_ref": "wp:D4-33!F12",
                            },
                        ],
                    }
                ],
            }
        ],
    }
    return parse_contract(payload, adapter_id=STATIC_CONTRACT_ID)


class TestSplitDefinedNameRef:
    def test_strips_quoted_sheet_and_dollar(self) -> None:
        sheet, a1 = _split_defined_name_ref("'其他业务毛利率分析表D433'!$E$12:$M$23")
        assert sheet == "其他业务毛利率分析表D433"
        assert a1 == "E12:M23"

    def test_strips_unquoted_sheet(self) -> None:
        sheet, a1 = _split_defined_name_ref("Sheet1!$A$1:$C$3")
        assert sheet == "Sheet1"
        assert a1 == "A1:C3"

    def test_missing_sheet_prefix_fail_closed(self) -> None:
        with pytest.raises(ManagedRegionResolutionError, match="缺 sheet 前缀"):
            _split_defined_name_ref("$A$1:$C$3")


class TestResolveStaticRegion:
    def _region(self, data: bytes):
        binding = ExcelIdentityBinding(
            table_key=STATIC_TABLE_KEY, defined_name=STATIC_DEFINED_NAME
        )
        with zipfile.ZipFile(io.BytesIO(data)) as zf:
            return _resolve_static_region(zf, contract=_static_contract(), binding=binding)

    def test_resolves_geometry_from_defined_name(self) -> None:
        region = self._region(_make_static_workbook())
        assert region.sheet_name == STATIC_SHEET
        assert region.first_column == "E"
        assert region.last_column == "M"
        assert region.first_row == 12
        assert region.last_row == 13
        assert region.uuid_column == ""  # Property 6：静态区无 UUID 列

    def test_missing_defined_name_is_carrier_missing(self) -> None:
        wb = openpyxl.Workbook()
        wb.active.title = STATIC_SHEET
        buf = io.BytesIO()
        wb.save(buf)
        with pytest.raises(IdentityCarrierMissingError, match="反读不到"):
            self._region(buf.getvalue())

    def test_duplicate_defined_name_fail_closed(self) -> None:
        with pytest.raises(ManagedRegionResolutionError, match="命中 2 个"):
            self._region(_make_static_workbook(duplicate=True))

    def test_defined_name_on_hidden_meta_sheet_rejected(self) -> None:
        with pytest.raises(ManagedRegionResolutionError, match="隐藏 metadata sheet"):
            self._region(_make_static_workbook(on_hidden_meta=True))

    def test_local_scope_defined_name_not_matched(self) -> None:
        # sheet-local scope 不是 workbook-scope → 反读不到 → carrier missing
        with pytest.raises(IdentityCarrierMissingError):
            self._region(_make_static_workbook(workbook_scope=False))


# ═══════════════════════════════════════════════════════════════════════════
# Task 3 · managed_tables_of 静态路径（Requirement 3.1 / 1.5 / Property 5）
# ═══════════════════════════════════════════════════════════════════════════

from app.services.workpaper_sync.excel_extract import managed_tables_of


def _contract_with(sheet_tables: list) -> object:
    import hashlib

    def _d(s: str) -> str:
        return hashlib.sha256(s.encode()).hexdigest()

    payload = {
        "schema_version": "contract-definition:v1",
        "contract_id": STATIC_CONTRACT_ID,
        "semantic_version": "1.0.0",
        "review_status": "reviewed",
        "document_type": "xlsx",
        "template_definition_sha256": _d("mt-template-definition"),
        "instrumentation_definition_sha256": _d("mt-instrumentation-definition"),
        "template": {
            "relative_path": "D/mixed.xlsx",
            "template_sha256": _d("mt-template"),
            "normalized_structure_hash": _d("mt-normalized-structure"),
        },
        "identity_carriers": ["defined_name", "excel_table", "hidden_uuid_column"],
        "sheets": [
            {
                "sheet_key": "mixed-managed",
                "excel_name": STATIC_SHEET,
                "locator": {"anchor": "defined_name_ref"},
                "tables": sheet_tables,
            }
        ],
    }
    return parse_contract(payload, adapter_id=STATIC_CONTRACT_ID)


def _static_table_spec(table_key: str = STATIC_TABLE_KEY) -> dict:
    return {
        "table_key": table_key,
        "anchor": "E12",
        "header_rows": 1,
        "fields": [
            {
                "stable_field_key": f"{table_key}/rev_m1",
                "json_pointer": "/months/0/revenue",
                "column_key": "rev_m1",
                "cell": {"column": "E", "row_from": 12},
                "mode": "editable",
                "value_type": "amount",
                "source_ref": "wp:D4-33!E12",
            },
            {
                "stable_field_key": f"{table_key}/cost_m1",
                "json_pointer": "/months/0/cost",
                "column_key": "cost_m1",
                "cell": {"column": "F", "row_from": 12},
                "mode": "editable",
                "value_type": "amount",
                "source_ref": "wp:D4-33!F12",
            },
        ],
    }


def _dynamic_table_spec(table_key: str = "dyn_rows") -> dict:
    return {
        "table_key": table_key,
        "anchor": "A20",
        "header_rows": 1,
        "row_identity": {"kind": "field", "json_pointer": "/rows/*/rowUuid"},
        "delete_policy": "tombstone",
        "fields": [
            {
                "stable_field_key": f"{table_key}/name",
                "json_pointer": "/rows/{row_uuid}/name",
                "column_key": "name",
                "cell": {"column": "A", "row_from": "row_identity"},
                "mode": "editable",
                "value_type": "text",
                "source_ref": "wp:D4-33!A20",
            }
        ],
    }


class TestManagedTablesStaticPath:
    def test_static_binding_returns_none_dynamic(self) -> None:
        contract = _contract_with([_static_table_spec()])
        binding = ExcelIdentityBinding(
            table_key=STATIC_TABLE_KEY, defined_name=STATIC_DEFINED_NAME
        )
        dynamic, statics = managed_tables_of(contract, binding=binding)
        assert dynamic is None  # 静态区无动态表
        assert len(statics) == 1
        assert statics[0].table_key == STATIC_TABLE_KEY

    def test_static_binding_pointing_at_dynamic_table_fail_closed(self) -> None:
        # 契约里 STATIC_TABLE_KEY 其实是动态表 → static binding 指向它 → 混填 fail-closed
        dyn = _dynamic_table_spec(table_key=STATIC_TABLE_KEY)
        contract = _contract_with([dyn])
        binding = ExcelIdentityBinding(
            table_key=STATIC_TABLE_KEY, defined_name=STATIC_DEFINED_NAME
        )
        with pytest.raises(ManagedRegionResolutionError, match="指向了动态表"):
            managed_tables_of(contract, binding=binding)

    def test_static_binding_sheet_without_static_table_fail_closed(self) -> None:
        # sheet 只有动态表、无静态表；binding.table_key 指向动态表 → 先撞「指向动态表」
        # 用一个不存在的 table_key 会先撞 sheet 定位（None）→ 契约不匹配
        dyn = _dynamic_table_spec(table_key="only_dyn")
        contract = _contract_with([dyn])
        binding = ExcelIdentityBinding(
            table_key="only_dyn", defined_name=STATIC_DEFINED_NAME
        )
        with pytest.raises(ManagedRegionResolutionError, match="指向了动态表"):
            managed_tables_of(contract, binding=binding)


class TestManagedTablesDynamicUnchanged:
    """动态 binding 的 has_dynamic_rows raise 逐条保留（Requirement 1.5 / Property 5 Task3 锚点）。

    Property 5 的核心：若把「按 uuid 空隐式分派」变异注入 kind，一个动态 binding（uuid 非空）
    仍判 dynamic（对），但把「静态 binding 误当动态」的边界——本测试钉住动态 binding 走
    managed_tables_of 时**必经** has_dynamic_rows raise（静态路径绕过它），二者行为可区分。
    """

    def test_dynamic_binding_returns_dynamic_and_statics(self) -> None:
        contract = _contract_with([_dynamic_table_spec(), _static_table_spec("footer_static")])
        binding = ExcelIdentityBinding(
            table_key="dyn_rows", table_name="GT_DYN", uuid_column="Z"
        )
        dynamic, statics = managed_tables_of(contract, binding=binding)
        assert dynamic is not None
        assert dynamic.table_key == "dyn_rows"
        assert {t.table_key for t in statics} == {"footer_static"}

    def test_dynamic_binding_pointing_at_static_table_fail_closed(self) -> None:
        # 动态 binding（uuid 非空）指向静态表 → 动态路径的 has_dynamic_rows raise 触发
        contract = _contract_with([_static_table_spec("s_only")])
        binding = ExcelIdentityBinding(
            table_key="s_only", table_name="GT_S", uuid_column="Z"
        )
        with pytest.raises(ManagedRegionResolutionError, match="没有 row_identity"):
            managed_tables_of(contract, binding=binding)


# ═══════════════════════════════════════════════════════════════════════════
# Task 4 · extract 静态链部件（Requirement 3.2/3.3 / Property 6）
# 注：完整 extract→materialize 往返（Property 1）在 Task 5 闭合（需 materialize + 注入
# definedName 的真实 substrate）。本段测 extract 侧可独立验证的部件：坐标集/列集/保留门。
# ═══════════════════════════════════════════════════════════════════════════

from app.services.workpaper_sync.excel_extract import (
    ManagedRegion,
    _assert_static_anchor_retained,
    _managed_coordinates,
    _needed_columns_and_rows,
)
from app.services.workpaper_sync.excel_entry_gate import parse_identity_inventory


def _static_region_obj() -> ManagedRegion:
    return ManagedRegion(
        table_key=STATIC_TABLE_KEY,
        table_name="",
        sheet_name=STATIC_SHEET,
        sheet_part="xl/worksheets/sheet1.xml",
        table_ref="E12:F12",
        first_row=12,
        last_row=12,
        first_column="E",
        last_column="F",
        uuid_column="",
    )


class TestManagedCoordinatesStatic:
    def test_no_ghost_uuid_coords(self) -> None:
        # Property 6：静态区受管坐标 = 静态 fields 绝对坐标集，不含任何 uuid_column{row}
        contract = _contract_with([_static_table_spec()])
        binding = ExcelIdentityBinding(
            table_key=STATIC_TABLE_KEY, defined_name=STATIC_DEFINED_NAME
        )
        coords = _managed_coordinates(
            contract=contract, region=_static_region_obj(), binding=binding, scan=None
        )
        assert coords == frozenset({"E12", "F12"})
        # 无空列坐标（uuid_column="" 会产生 "12" 之类幽灵坐标）
        assert not any(c.lstrip("ABCDEFGHIJKLMNOPQRSTUVWXYZ") == c for c in coords)


class TestNeededColumnsStatic:
    def test_columns_exclude_uuid(self) -> None:
        contract = _contract_with([_static_table_spec()])
        binding = ExcelIdentityBinding(
            table_key=STATIC_TABLE_KEY, defined_name=STATIC_DEFINED_NAME
        )
        columns, rows = _needed_columns_and_rows(
            contract=contract, region=_static_region_obj(), binding=binding
        )
        assert columns == frozenset({"E", "F"})  # 不含 uuid_column ""
        assert "" not in columns
        assert 12 in rows


def _static_inventory_raw(defined_names: list[str]) -> dict:
    return {
        "hidden_sheet": {
            "present": True,
            "is_hidden": True,
            "excluded_from_business_enumeration": True,
        },
        "defined_name": {"names": {n: f"'{STATIC_SHEET}'!$E$12:$M$23" for n in defined_names}},
        "excel_table": {"present": False, "table_ref": ""},
        "hidden_uuid_column": {
            "resolved_sheet_by": None,
            "uuid_column_hidden": False,
            "row_uuids": {},
            "duplicate_row_uuids": [],
            "empty_row_uuids": [],
        },
    }


def _runtime_static_inventory(defined_names: tuple[str, ...], *, hidden_present: bool = True):
    from app.services.workpaper_sync.excel_extract import RuntimeIdentityInventory
    from app.services.workpaper_sync.excel_extract import TABLE_SHEET_ANCHOR  # noqa: F401

    return RuntimeIdentityInventory(
        hidden_sheet_present=hidden_present,
        hidden_sheet_is_hidden=True,
        excluded_from_business_enumeration=True,
        defined_names=defined_names,
        table_present=True,
        table_ref="E12:F12",
        table_sheet=STATIC_SHEET,
        resolved_sheet_by=None,
        uuid_column="",
        uuid_column_hidden=False,
        row_uuids={},
        duplicate_row_uuids=(),
        empty_row_uuids=(),
        business_sheets=(STATIC_SHEET,),
    )


class TestStaticAnchorRetained:
    def test_anchor_present_passes(self) -> None:
        expected = parse_identity_inventory(_static_inventory_raw([STATIC_DEFINED_NAME]))
        observed = _runtime_static_inventory((STATIC_DEFINED_NAME,))
        binding = ExcelIdentityBinding(
            table_key=STATIC_TABLE_KEY, defined_name=STATIC_DEFINED_NAME
        )
        # 不抛即通过
        _assert_static_anchor_retained(
            expected=expected, observed=observed, binding=binding, entry_id="e1"
        )

    def test_anchor_lost_in_observed_fail_closed(self) -> None:
        from app.services.workpaper_sync.excel_extract import IdentityRetentionError

        expected = parse_identity_inventory(_static_inventory_raw([STATIC_DEFINED_NAME]))
        observed = _runtime_static_inventory(())  # OO 往返后 definedName 丢了
        binding = ExcelIdentityBinding(
            table_key=STATIC_TABLE_KEY, defined_name=STATIC_DEFINED_NAME
        )
        with pytest.raises(IdentityRetentionError, match="在 OO 往返后丢失"):
            _assert_static_anchor_retained(
                expected=expected, observed=observed, binding=binding, entry_id="e1"
            )

    def test_hidden_sheet_lost_fail_closed(self) -> None:
        from app.services.workpaper_sync.excel_extract import IdentityRetentionError

        expected = parse_identity_inventory(_static_inventory_raw([STATIC_DEFINED_NAME]))
        observed = _runtime_static_inventory((STATIC_DEFINED_NAME,), hidden_present=False)
        binding = ExcelIdentityBinding(
            table_key=STATIC_TABLE_KEY, defined_name=STATIC_DEFINED_NAME
        )
        with pytest.raises(IdentityRetentionError, match="隐藏 metadata sheet"):
            _assert_static_anchor_retained(
                expected=expected, observed=observed, binding=binding, entry_id="e1"
            )


# ═══════════════════════════════════════════════════════════════════════════
# Task 5 · materialize 静态链 + 闭合往返（Property 1）
# ═══════════════════════════════════════════════════════════════════════════

from app.services.workpaper_sync import excel_materialize as M
from app.services.workpaper_sync.adapters.base import (
    FieldValue,
    Projection,
    SubstrateRole,
)
from app.services.workpaper_sync.excel_entry_gate import (
    AdapterBuild,
    FrozenEntryDefinitions,
)
from app.services.workpaper_sync.models import ArtifactKind, ArtifactState


def _read_entries(data: bytes) -> dict:
    with zipfile.ZipFile(io.BytesIO(data)) as zf:
        return {name: zf.read(name) for name in zf.namelist()}


def _make_full_static_substrate(values: dict[str, float]) -> bytes:
    """造完整静态 substrate：受管静态 sheet + definedName + 隐藏 _GT_SYNC sheet。

    受管区 E12:F12（rev/cost 两输入），G12 = 一个公式 cell（=E12-F12，formula_mask）。
    """
    from app.services.excel_structure_fingerprint import GT_SYNC_SHEET_NAME

    wb = openpyxl.Workbook()
    ws = wb.active
    ws.title = STATIC_SHEET
    ws["E12"] = values.get("E12", 0)
    ws["F12"] = values.get("F12", 0)
    ws["G12"] = "=E12-F12"  # 区内公式（formula_mask 保护）
    meta = wb.create_sheet(GT_SYNC_SHEET_NAME)
    meta["A1"] = "GT_SYNC_SCHEMA_VERSION"
    meta["B1"] = "1"
    meta.sheet_state = "hidden"
    wb.defined_names[STATIC_DEFINED_NAME] = DefinedName(
        STATIC_DEFINED_NAME, attr_text=f"'{STATIC_SHEET}'!$E$12:$G$12"
    )
    buf = io.BytesIO()
    wb.save(buf)
    return buf.getvalue()


def _static_contract_with_formula() -> object:
    import hashlib

    def _d(s: str) -> str:
        return hashlib.sha256(s.encode()).hexdigest()

    payload = {
        "schema_version": "contract-definition:v1",
        "contract_id": STATIC_CONTRACT_ID,
        "semantic_version": "1.0.0",
        "review_status": "reviewed",
        "document_type": "xlsx",
        "template_definition_sha256": _d("d433f-template-definition"),
        "instrumentation_definition_sha256": _d("d433f-instrumentation-definition"),
        "template": {
            "relative_path": "D/d433f.xlsx",
            "template_sha256": _d("d433f-template"),
            "normalized_structure_hash": _d("d433f-normalized-structure"),
        },
        "identity_carriers": ["defined_name"],
        "sheets": [
            {
                "sheet_key": "d433-managed",
                "excel_name": STATIC_SHEET,
                "locator": {"anchor": "defined_name_ref"},
                "tables": [
                    {
                        "table_key": STATIC_TABLE_KEY,
                        "anchor": "E12",
                        "header_rows": 1,
                        "formula_mask": ["G12:G12"],
                        "fields": [
                            {
                                "stable_field_key": f"{STATIC_TABLE_KEY}/rev_m1",
                                "json_pointer": "/months/0/revenue",
                                "column_key": "rev_m1",
                                "cell": {"column": "E", "row_from": 12},
                                "mode": "editable",
                                "value_type": "amount",
                                "source_ref": "wp:D4-33!E12",
                            },
                            {
                                "stable_field_key": f"{STATIC_TABLE_KEY}/cost_m1",
                                "json_pointer": "/months/0/cost",
                                "column_key": "cost_m1",
                                "cell": {"column": "F", "row_from": 12},
                                "mode": "editable",
                                "value_type": "amount",
                                "source_ref": "wp:D4-33!F12",
                            },
                            {
                                "stable_field_key": f"{STATIC_TABLE_KEY}/margin_m1",
                                "json_pointer": "/months/0/margin",
                                "column_key": "margin_m1",
                                "cell": {"column": "G", "row_from": 12},
                                "mode": "formula",
                                "value_type": "amount",
                                "source_ref": "wp:D4-33!G12",
                            },
                        ],
                    }
                ],
            }
        ],
    }
    return parse_contract(payload, adapter_id=STATIC_CONTRACT_ID)


def _static_definitions(contract: object) -> FrozenEntryDefinitions:
    import hashlib
    import uuid as _uuid

    from app.services.workpaper_sync.models import (
        AuthorityModel,
        BundleSlot,
        BundleSlotSpec,
        DefinitionState,
    )
    from app.services.workpaper_sync.resolution import DefinitionBundleSnapshot

    def _d(s: str) -> str:
        return hashlib.sha256(s.encode()).hexdigest()

    def slot(kind, digest):
        return BundleSlotSpec(kind, "definition", f"definition:{_uuid.uuid4()}", digest)

    bundle = DefinitionBundleSnapshot(
        bundle_id=_uuid.uuid4(),
        bundle_sha256=_d("d433-bundle"),
        schema_version="definition-bundle:v1",
        state=DefinitionState.approved,
        authority_model=AuthorityModel.projection_contract,
        authority_model_definition_id=_uuid.uuid4(),
        authority_model_definition_sha256=_d("d433-authority"),
        slots={
            BundleSlot.template: slot(BundleSlot.template, contract.template_definition_sha256),
            BundleSlot.instrumentation: slot(
                BundleSlot.instrumentation, contract.instrumentation_definition_sha256
            ),
            BundleSlot.contract: slot(BundleSlot.contract, contract.canonical_sha256),
        },
    )
    return FrozenEntryDefinitions(
        entry_id="xlsx/gt-d4-33-static-test",
        bundle=bundle,
        contract=contract,
        adapter_build=AdapterBuild(
            adapter_id=STATIC_CONTRACT_ID,
            adapter_build_digest=_d("d433-adapter-build"),
            document_type="xlsx",
            contract_version="1.0.0",
        ),
        identity_inventory=parse_identity_inventory(
            _static_inventory_raw([STATIC_DEFINED_NAME])
        ),
        business_sheets=(),
        dynamic_column_keys={},
        structure_inventory_size=0,
    )


def _static_binding() -> ExcelIdentityBinding:
    return ExcelIdentityBinding(
        table_key=STATIC_TABLE_KEY, defined_name=STATIC_DEFINED_NAME
    )


def _projection_for(contract: object, *, rev: float, cost: float) -> Projection:
    def fv(key: str, value, vtype, mode):
        spec = contract.field_by_stable_key(key)
        return FieldValue(key, value, spec.value_type, spec.mode)

    values = {
        f"{STATIC_TABLE_KEY}/rev_m1": fv(
            f"{STATIC_TABLE_KEY}/rev_m1", rev, "amount", "editable"
        ),
        f"{STATIC_TABLE_KEY}/cost_m1": fv(
            f"{STATIC_TABLE_KEY}/cost_m1", cost, "amount", "editable"
        ),
        # formula 字段的缓存值（派生），plan 应保留其 <f> 不当普通值写
        f"{STATIC_TABLE_KEY}/margin_m1": fv(
            f"{STATIC_TABLE_KEY}/margin_m1", rev - cost, "amount", "formula"
        ),
    }
    return Projection(
        contract_id=contract.contract_id,
        semantic_version=contract.semantic_version,
        document_type=contract.document_type,
        values=values,
        row_keys={},
    )


class TestStaticRoundtrip:
    """Property 1：静态区 materialize→extract 往返受管值逐字段相等；公式归 formula_inventory。"""

    def test_write_then_read_equal(self, tmp_path) -> None:
        from app.services.workpaper_sync.excel_extract import extract_projection

        contract = _static_contract_with_formula()
        definitions = _static_definitions(contract)
        binding = _static_binding()
        substrate = _make_full_static_substrate({"E12": 111.0, "F12": 22.0})

        # 1) plan 写新值 rev=888 cost=100
        region_path = tmp_path / "base.xlsx"
        region_path.write_bytes(substrate)
        with zipfile.ZipFile(io.BytesIO(substrate)) as zf:
            region = M.resolve_managed_region(zf, contract=contract, binding=binding)
        plan = M.plan_managed_writes(
            projection=_projection_for(contract, rev=888.0, cost=100.0),
            contract=contract,
            binding=binding,
            region=region,
            scan=None,
            substrate_entries=_read_entries(substrate),
            substrate_formulas={},
            runtime_binding={},
        )
        assert plan.row_shift is None
        assert plan.footer_marker_row is None
        # G12 是 formula → 应进 preserved_formulas
        assert f"{STATIC_TABLE_KEY}/margin_m1" in plan.preserved_formulas
        # G12 的写入是 cached_value_only（保留 <f>），E12/F12 是普通值写
        kind_by_coord = {w.coord: w.kind for w in plan.writes}
        assert kind_by_coord["G12"] is M.CellWriteKind.cached_value_only
        assert kind_by_coord["E12"] is M.CellWriteKind.number_literal
        assert kind_by_coord["F12"] is M.CellWriteKind.number_literal
        # 静态区无 identity 写入（无 UUID 列）
        assert plan.identity_writes == ()

        # 2) apply → staged
        staged = M.apply_plan_zip(substrate, plan)
        staged_path = tmp_path / "staged.xlsx"
        staged_path.write_bytes(staged)

        # 3) extract 反读
        outcome = extract_projection(
            artifact=staged_path,
            definitions=definitions,
            binding=binding,
            substrate_role=SubstrateRole.published_representation,
            artifact_kind=ArtifactKind.canonical,
            artifact_state=ArtifactState.published,
        )
        rev_back = outcome.projection.get(f"{STATIC_TABLE_KEY}/rev_m1")
        cost_back = outcome.projection.get(f"{STATIC_TABLE_KEY}/cost_m1")
        assert float(rev_back.value) == 888.0
        assert float(cost_back.value) == 100.0
        # 公式 cell 归 formula_inventory，不当受管值
        assert f"{STATIC_TABLE_KEY}/margin_m1" in outcome.formula_inventory


# ═══════════════════════════════════════════════════════════════════════════
# Task 6 · observer 泛化（Requirement 5.1/5.3 / Property 3/7）
# ═══════════════════════════════════════════════════════════════════════════

from app.services.workpaper_sync.published_identity_observer import (
    _collect_static_region_physical,
    _frozen_sheet_anchors,
    collect_workbook_structure,
)


class TestFrozenSheetAnchorsStatic:
    def test_static_payload_yields_region_kind_static(self) -> None:
        payload = {
            "managed_sheets": [
                {
                    "sheet_key": "d433-managed",
                    "region_boundary_locator": {
                        "anchor": "defined_name_ref",
                        "defined_name": STATIC_DEFINED_NAME,
                        "region_kind": "static",
                    },
                }
            ]
        }
        anchors = _frozen_sheet_anchors(payload)
        assert len(anchors) == 1
        assert anchors[0]["anchor"] == "defined_name_ref"
        assert anchors[0]["region_kind"] == "static"
        assert anchors[0]["defined_name"] == STATIC_DEFINED_NAME

    def test_transposed_payload_has_empty_region_kind(self) -> None:
        # D4-29 payload 无 region_kind → 缺省 ""（走 transposed，零回归）
        payload = {
            "transposed_sheets": [
                {
                    "sheet_key": "d4-29",
                    "region_boundary_locator": {
                        "anchor": "defined_name_ref",
                        "defined_name": "GT_MANAGED_REGION_D429",
                    },
                }
            ]
        }
        anchors = _frozen_sheet_anchors(payload)
        assert len(anchors) == 1
        assert anchors[0]["region_kind"] == ""


class TestCollectStaticRegionPhysical:
    def test_resolves_sheet_name(self) -> None:
        data = _make_full_static_substrate({"E12": 1.0, "F12": 2.0})
        contract = _static_contract_with_formula()
        physical = _collect_static_region_physical(
            data=data, contract=contract, key="d433-managed",
            defined_name=STATIC_DEFINED_NAME,
        )
        assert physical == STATIC_SHEET

    def test_missing_defined_name_fail_closed(self) -> None:
        wb = openpyxl.Workbook()
        wb.active.title = STATIC_SHEET
        buf = io.BytesIO()
        wb.save(buf)
        with pytest.raises(ValueError, match="反读不到"):
            _collect_static_region_physical(
                data=buf.getvalue(), contract=_static_contract_with_formula(),
                key="d433-managed", defined_name=STATIC_DEFINED_NAME,
            )


class TestCollectWorkbookStructureStatic:
    def test_static_anchor_goes_through_static_path(self) -> None:
        # collect_workbook_structure 对 static anchor 走 _collect_static_region_physical，
        # 不调 extract_transposed_workbook（那是转置语义）。
        data = _make_full_static_substrate({"E12": 10.0, "F12": 3.0})
        contract = _static_contract_with_formula()
        anchors = [
            {
                "sheet_key": "d433-managed",
                "anchor": "defined_name_ref",
                "defined_name": STATIC_DEFINED_NAME,
                "region_kind": "static",
            }
        ]
        fingerprint, physical, primary_inventory, structure = collect_workbook_structure(
            data=data, contract=contract, sheet_anchors=anchors
        )
        assert physical["d433-managed"] == STATIC_SHEET
        # 静态区无 primary_inventory（那是动态 row UUID 路径产的）
        assert primary_inventory is None
        # 结构清册含静态 fields 的绝对坐标（E12/F12/G12 由 _cell_coordinates_for 静态路径产）
        locators = {loc for (_sk, _tk, _fk, loc) in structure}
        assert any(":12" in loc for loc in locators)


# ═══════════════════════════════════════════════════════════════════════════
# Task 7 · instrumentation 注入静态 definedName（Requirement 5.4/5.5）
# ═══════════════════════════════════════════════════════════════════════════

import pathlib

from app.services.workpaper_sync import excel_instrumentation as EI

_WP_TEMPLATES = pathlib.Path(__file__).resolve().parents[2] / "wp_templates"
_K11_TEMPLATE = _WP_TEMPLATES / "K" / "K11 资产减值损失.xlsx"


def _spec_with_static(static_sheet_name: str, defined_name: str, region_range: str):
    """动态 primary spec（K11）+ 一个静态受管区寄生声明。"""
    return EI.ExcelInstrumentationSpec(
        entry_id="k11.adjudication",
        template_id="K11",
        template_relative_path="K/K11 资产减值损失.xlsx",
        managed_sheet="审定表K11-1",
        first_data_row=7,
        last_data_row=25,
        footer_row=26,
        managed_last_col="L",
        uuid_col="N",
        table_name="GT_K11_1_ROWS",
        static_sheets=(
            {
                "sheet_key": "d433-managed",
                "excel_name": static_sheet_name,
                "region_boundary_locator": {
                    "anchor": "defined_name_ref",
                    "defined_name": defined_name,
                    "range": region_range,
                    "region_kind": "static",
                },
                "tables": [{"table_key": "d433-static"}],
            },
        ),
    )


def _first_other_sheet_name(data: bytes) -> str:
    import re as _re

    wb = _read_entries(data)["xl/workbook.xml"].decode("utf-8")
    names = _re.findall(r'<sheet [^>]*name="([^"]+)"', wb)
    for n in names:
        if n != "审定表K11-1":
            return n
    raise AssertionError("K11 模板只有一张 sheet？")


class TestStaticInstrumentation:
    @pytest.fixture(scope="class")
    def source(self) -> bytes:
        if not _K11_TEMPLATE.exists():
            pytest.skip(f"模板缺失: {_K11_TEMPLATE}")
        return _K11_TEMPLATE.read_bytes()

    def test_injects_workbook_scope_defined_name_no_table_no_uuid(self, source: bytes) -> None:
        static_sheet = _first_other_sheet_name(source)
        spec = _spec_with_static(static_sheet, "GT_MANAGED_REGION_D433T", "$E$12:$G$13")
        gate = EI.ExcelIdentityCarrierGate.load()
        result = EI.instrument_workbook_bytes_multi(source, [spec], gate=gate)
        entries = _read_entries(result.instrumented_bytes)
        wb = entries["xl/workbook.xml"].decode("utf-8")
        # (a) 注入了 workbook-scope definedName（无 localSheetId）
        assert "GT_MANAGED_REGION_D433T" in wb
        import re as _re

        m = _re.search(
            r'<definedName name="GT_MANAGED_REGION_D433T"([^>]*)>([^<]+)</definedName>', wb
        )
        assert m is not None
        assert "localSheetId" not in m.group(1)  # workbook-scope
        assert "$E$12:$G$13" in m.group(2)
        # definedName_refs 里也有它
        assert "GT_MANAGED_REGION_D433T" in result.defined_name_refs

    def test_static_defined_name_resolvable_by_extract(self, source: bytes) -> None:
        # 注入后用 extract 的 _resolve_static_region 能反解出 region（端到端锚点闭合）
        static_sheet = _first_other_sheet_name(source)
        spec = _spec_with_static(static_sheet, "GT_MANAGED_REGION_D433T", "$E$12:$G$13")
        gate = EI.ExcelIdentityCarrierGate.load()
        result = EI.instrument_workbook_bytes_multi(source, [spec], gate=gate)
        from app.services.workpaper_sync.excel_extract import (
            _resolve_static_region,
        )

        binding = ExcelIdentityBinding(
            table_key="d433-static", defined_name="GT_MANAGED_REGION_D433T"
        )
        with zipfile.ZipFile(io.BytesIO(result.instrumented_bytes)) as zf:
            region = _resolve_static_region(
                zf, contract=_static_contract_with_formula(), binding=binding
            )
        assert region.sheet_name == static_sheet
        assert region.first_column == "E"
        assert region.last_column == "G"
        assert region.first_row == 12
        assert region.uuid_column == ""

    def test_frozen_anchors_from_payload_yield_static_kind(self, source: bytes) -> None:
        spec = _spec_with_static(
            _first_other_sheet_name(source), "GT_MANAGED_REGION_D433T", "$E$12:$G$13"
        )
        import hashlib as _h

        payload = EI.build_instrumentation_payload_for_sheets(
            specs=[spec],
            template_definition_sha256=_h.sha256(b"td").hexdigest(),
            template_sha256=_h.sha256(b"ts").hexdigest(),
            gate=EI.ExcelIdentityCarrierGate.load(),
        )
        assert "static_sheets" in payload
        anchors = _frozen_sheet_anchors(payload)
        static_anchors = [a for a in anchors if a.get("region_kind") == "static"]
        assert len(static_anchors) == 1
        assert static_anchors[0]["defined_name"] == "GT_MANAGED_REGION_D433T"

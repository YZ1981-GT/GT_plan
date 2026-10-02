"""D4-4 Task 17：真实 materialize 扩容与 footer/身份持久化。

这条用例必须走真实 D4 模板、真实 instrumentation、真实 identity inventory 和
ExcelSyncAdapter.materialize。只构造 D4-4 projection 会漏掉同一 entry 的 sibling
binding，因此 baseline 先由完整 adapter.extract() 反读，再仅覆盖 D4-4。
"""
from __future__ import annotations

import hashlib
import os
import sys
from pathlib import Path
from typing import Any, Mapping
from uuid import NAMESPACE_URL, uuid5

import openpyxl

_REPO = Path(__file__).resolve().parents[3]
_BACKEND = _REPO / "backend"
if str(_BACKEND) not in sys.path:  # pragma: no cover - import 环境自举
    sys.path.insert(0, str(_BACKEND))
os.environ.setdefault("DB_DISABLE_SSL", "True")

from app.services.excel_structure_fingerprint import (  # noqa: E402
    GT_SYNC_SHEET_NAME,
    identity_inventory,
)
from app.services.workpaper_sync import excel_instrumentation as EI  # noqa: E402
from app.services.workpaper_sync import phase5_d4_adjustment_sheet as D44  # noqa: E402
from app.services.workpaper_sync import phase5_d4_revenue_detail as D4  # noqa: E402
from app.services.workpaper_sync.adapters import excel as AX  # noqa: E402
from app.services.workpaper_sync.contracts import parse_contract  # noqa: E402
from app.services.workpaper_sync.excel_entry_gate import (  # noqa: E402
    AdapterBuild,
    FrozenEntryDefinitions,
    parse_identity_inventory,
)
from app.services.workpaper_sync.excel_extract import ExcelIdentityBinding  # noqa: E402
from app.services.workpaper_sync.projection_first_publication import (  # noqa: E402
    _align_specs_to_sibling_tables,
    _static_region_bindings,
    overlay_store_on_baseline_projection,
)
from tests.workpaper_sync.d4_materialize_harness import (  # noqa: E402
    MaterializeCallCounter,
    _bundle,
    _text_safe,
)


def _existing_d44_ids(inventory: Mapping[str, Any]) -> tuple[str, ...]:
    row_uuids = inventory["hidden_uuid_column"]["row_uuids"]
    ids: list[str] = []
    for row in range(D44.FIRST_DATA_ROW_D44, D44.LAST_DATA_ROW_D44 + 1):
        value = row_uuids.get(row) or row_uuids.get(str(row))
        assert value, f"D4-4 行 {row} 缺少 instrumentation UUID: {row_uuids!r}"
        ids.append(str(value))
    assert len(ids) == 15 and len(set(ids)) == 15
    return tuple(ids)


def _d44_rows(row_ids: tuple[str, ...]) -> list[dict[str, Any]]:
    rows: list[dict[str, Any]] = []
    for ordinal, row_id in enumerate(row_ids, start=1):
        rows.append(
            {
                D44.ROW_IDENTITY_KEY_D44: row_id,
                "description": f"Task 17 调整事项 {ordinal}",
                "category": "账项调整",
                "reportItem": f"营业收入项目 {ordinal}",
                "accountName": f"测试科目 {ordinal}",
                "noteItem": f"附注项目 {ordinal}",
                "placeholder": "",
                "debitAmount": float(ordinal * 100),
                "creditAmount": float(ordinal * 10),
                "indexRef": f"D4-4-{ordinal:02d}",
                "remark": f"真实 materialize 扩容行 {ordinal}",
            }
        )
    return rows


def _d44_primary_and_siblings(
    *, contract: Any, d44_spec: Any
) -> tuple[ExcelIdentityBinding, tuple[ExcelIdentityBinding, ...]]:
    primary = ExcelIdentityBinding(
        table_name=str(d44_spec.table_name),
        uuid_column=str(d44_spec.uuid_col),
        table_key=D44.TABLE_KEY_D44,
        metadata_sheet=GT_SYNC_SHEET_NAME,
        defined_name_prefix=str(getattr(d44_spec, "defined_name_prefix", None) or "GT_"),
        tombstoned_row_keys=(),
        dynamic_column_columns={},
    )
    siblings = [
        ExcelIdentityBinding(
            table_name=str(spec.table_name),
            uuid_column=str(spec.uuid_col),
            table_key=str(dynamic.table_key),
            metadata_sheet=GT_SYNC_SHEET_NAME,
            defined_name_prefix=str(
                getattr(spec, "defined_name_prefix", None) or "GT_"
            ),
            tombstoned_row_keys=(),
            dynamic_column_columns={},
        )
        for spec, dynamic in _align_specs_to_sibling_tables(
            provider=D4, contract=contract, primary=primary
        )
    ]
    siblings.extend(
        _static_region_bindings(provider=D4, metadata_sheet=GT_SYNC_SHEET_NAME)
    )
    return primary, tuple(siblings)


def _definitions(
    *, contract: Any, inventory: Mapping[str, Any]
) -> FrozenEntryDefinitions:
    return FrozenEntryDefinitions(
        entry_id=D4.ENTRY_ID,
        bundle=_bundle(contract),
        contract=contract,
        adapter_build=AdapterBuild(
            adapter_id=D4.ADAPTER_ID,
            adapter_build_digest=hashlib.sha256(
                b"d4-4-task-17-materialize-adapter"
            ).hexdigest(),
            document_type="xlsx",
            contract_version=contract.semantic_version,
        ),
        identity_inventory=parse_identity_inventory(inventory),
        business_sheets=(),
        dynamic_column_keys={},
        structure_inventory_size=0,
    )


def test_d44_materialize_expands_fifteen_rows_and_keeps_footer_and_identities(
    tmp_path: Path,
) -> None:
    """15 个已有身份 + 5 个新身份必须真实落成 20 行，并保留完整 footer marker。"""
    D4.assert_contract_file_matches_source()
    contract = parse_contract(D4.build_contract_payload(), adapter_id=D4.ADAPTER_ID)
    instrumented = EI.instrument_workbook_bytes_multi(
        D4.read_authoritative_template(),
        D4.instrumentation_specs(),
        gate=D4.excel_carrier_gate(),
    )
    d44_specs = [
        spec
        for spec in D4.instrumentation_specs()
        if str(spec.managed_sheet) == D44.MANAGED_SHEET_D44
    ]
    assert len(d44_specs) == 1, f"D4-4 instrumentation spec 数量异常: {d44_specs!r}"
    d44_spec = d44_specs[0]
    assert str(d44_spec.uuid_col) == D44.UUID_COL_D44

    inventory = identity_inventory(
        instrumented.instrumented_bytes,
        expected_table=str(d44_spec.table_name),
        uuid_column_letter=D44.UUID_COL_D44,
    )
    existing_ids = _existing_d44_ids(inventory)
    new_ids = tuple(
        str(uuid5(NAMESPACE_URL, f"gt-plan/d4-4/task17/new-row/{index}"))
        for index in range(1, 6)
    )
    target_ids = existing_ids + new_ids

    primary, siblings = _d44_primary_and_siblings(
        contract=contract, d44_spec=d44_spec
    )
    adapter = AX.build_excel_adapter(
        definitions=_definitions(contract=contract, inventory=inventory),
        binding=primary,
        sibling_bindings=siblings,
        direction="html_to_oo",
    )
    all_bindings = adapter._all_bindings()
    assert all_bindings[0].table_key == D44.TABLE_KEY_D44
    assert len({binding.table_key for binding in all_bindings}) == len(all_bindings)
    binding_count = len(all_bindings)
    assert binding_count == 1 + len(siblings)

    base = tmp_path / "d4-4-task17-base.xlsx"
    base.write_bytes(instrumented.instrumented_bytes)
    baseline = _text_safe(adapter.extract(artifact=base, contract=contract))
    d44_projection = D44.build_store_projection_d44(
        _d44_rows(target_ids), contract=contract
    )
    projection = overlay_store_on_baseline_projection(
        baseline=baseline, store_projection=d44_projection
    )
    assert projection.row_keys[D44.TABLE_KEY_D44] == target_ids

    output = tmp_path / AX.STAGING_NAMESPACE / "d4-4-task17-capacity.xlsx"
    output.parent.mkdir(parents=True, exist_ok=True)
    counter = MaterializeCallCounter()
    with counter.installed():
        result = adapter.materialize(
            substrate=base,
            projection=projection,
            output=output,
            contract=contract,
        )

    assert output.is_file(), "materialize 未生成 staged 文件"
    assert output.read_bytes()[:2] == b"PK", "staged 产物不是可解析的 xlsx zip"
    assert result.row_shift is not None, "D4-4 扩容未产生 row shift"
    assert result.row_shift.table_key == D44.TABLE_KEY_D44
    assert result.row_shift.count == 5
    assert result.per_table_shift is not None
    assert result.per_table_shift[D44.TABLE_KEY_D44][0].count == 5
    assert counter.trip_count == binding_count, counter.as_dict()
    assert counter.single_pass_trip_count == 0, counter.as_dict()

    staged_inventory = identity_inventory(
        output.read_bytes(),
        expected_table=str(d44_spec.table_name),
        uuid_column_letter=D44.UUID_COL_D44,
    )
    staged_ids = tuple(
        str(value)
        for value in staged_inventory["hidden_uuid_column"]["row_uuids"].values()
        if value
    )
    assert len(staged_ids) == 20
    assert len(set(staged_ids)) == 20
    assert set(staged_ids) == set(target_ids)

    workbook = openpyxl.load_workbook(output, data_only=False)
    try:
        worksheet = workbook[D44.MANAGED_SHEET_D44]
        marker_rows = [
            row
            for row in range(1, worksheet.max_row + 1)
            if str(worksheet.cell(row=row, column=1).value or "").strip()
            == D44.FOOTER_MARKER_D44
        ]
        assert len(marker_rows) == 1, (
            f"完整 footer marker 必须恰出现一次，实际行号 {marker_rows!r}"
        )
        marker_row = marker_rows[0]
        identity_rows = [
            row
            for row in range(D44.FIRST_DATA_ROW_D44, marker_row)
            if worksheet.cell(row=row, column=11).value
        ]
        assert identity_rows == list(
            range(D44.FIRST_DATA_ROW_D44, D44.FIRST_DATA_ROW_D44 + 20)
        )
        assert max(identity_rows) < marker_row
        k_values = [
            str(worksheet.cell(row=row, column=11).value) for row in identity_rows
        ]
        assert len(k_values) == 20
        assert len(set(k_values)) == 20
        assert set(k_values) == set(target_ids)
        assert sum(
            1
            for row in range(1, worksheet.max_row + 1)
            if str(worksheet.cell(row=row, column=1).value or "").strip()
            == D44.FOOTER_MARKER_D44
        ) == 1
    finally:
        workbook.close()

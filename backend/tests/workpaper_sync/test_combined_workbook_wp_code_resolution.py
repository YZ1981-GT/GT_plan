"""合册命名的 wp_code 解析（Lane B）。

spec: workpaper-sync-pure-static-lane-and-combined-workbook-resolution

═══ 本文件守什么 ═══

`backend/wp_templates/A/` 下两本「一个文件装两个 wp_code」的合册，其第二个码在
`wp_template_finder` 里解析不到自己所在的册：`find_template_file_any` 返回 `None`
（可见失败），`find_template_file` 返回**错的册**（父程序表，静默错答，更坏）。

新逻辑以**加法**方式插在两条链的精确位置，故：
* 当前能解析的码走向**一律不变**（对照组判据）；
* 只有当前失败的码进入新步骤（全域前后差集判据）。

判据纪律：分母一律现算；范围式命名 fail-closed 配**双向**变异证明；
收窄条件 `len(declared) < 2` 配**反向**变异证明。
"""
from __future__ import annotations

import json
import os
import subprocess
import sys
from pathlib import Path

import pytest

_BACKEND = Path(__file__).resolve().parents[2]
if str(_BACKEND) not in sys.path:
    sys.path.insert(0, str(_BACKEND))
os.environ.setdefault("DB_DISABLE_SSL", "True")

from app.services import wp_template_finder as F  # noqa: E402

from hypothesis import HealthCheck, given, settings  # noqa: E402
from hypothesis import strategies as st  # noqa: E402

#: 项目铁律：`max_examples=5`，禁 hypothesis 默认 100。
PBT = settings(
    max_examples=5, deadline=None,
    suppress_health_check=[HealthCheck.function_scoped_fixture, HealthCheck.too_slow],
)

COMBINED_A38 = "A/A3-7内部往来核对表、A3-8商誉减值测试.xlsx"
COMBINED_A22 = "A/A2-1、A2-2单体报表试算.xlsx"


def rel(p: Path | None) -> str | None:
    if p is None:
        return None
    return str(Path(p).resolve().relative_to(F.TEMPLATES_DIR.resolve())).replace("\\", "/")


def both_routes(wp_code: str) -> tuple[str | None, str | None]:
    """两路**权威侧**入口（绕开覆盖层薄封装，判据锁的是权威解析）。"""
    return (
        rel(F.find_template_file_unresolved(wp_code)),
        rel(F.find_template_file_any_unresolved(wp_code)),
    )


@pytest.fixture(scope="module")
def all_workbook_names() -> list[str]:
    """权威目录下全部 `.xlsx` / `.xlsm` 文件名（现算分母，禁写死）。"""
    names = [
        f.name
        for sub in F.TEMPLATES_DIR.iterdir() if sub.is_dir()
        for f in sub.iterdir() if f.suffix.lower() in (".xlsx", ".xlsm")
    ]
    assert names, "权威目录分母为空 —— 判据失去分母"
    return names


@pytest.fixture(scope="module")
def combined_books(all_workbook_names: list[str]) -> list[str]:
    """fail-closed 之后仍算合册的文件名（现算，禁写死成员）。"""
    hits = [n for n in all_workbook_names
            if len(F._literal_wp_codes_in_filename(n)) >= 2]
    assert hits, "合册分母为空 —— 本文件全部判据失去分母"
    return sorted(hits)


class TestModulePrivateSurface:
    """Requirement 10.1：五个符号都是**模块私有**，不新增公开入口。"""

    def test_five_private_symbols_exist(self) -> None:
        for name in ("_RANGE_MARKERS", "_FILENAME_CODE_RE",
                     "_literal_wp_codes_in_filename", "_combined_book_covers",
                     "_find_combined_workbook_declaring"):
            assert hasattr(F, name), f"缺模块私有符号 {name}"
            assert name.startswith("_"), f"{name} 不是模块私有"

    def test_no_new_public_entry_point(self) -> None:
        """不进模块末尾的覆盖层薄封装、不新增公开入口（Requirement 11.6）。"""
        public = {n for n in vars(F) if not n.startswith("_") and callable(vars(F)[n])}
        assert "find_combined_workbook_declaring" not in public
        # 覆盖层重绑定与 `*_unresolved` 别名保持不变
        for name in ("find_template_file_unresolved", "find_all_template_files_unresolved",
                     "find_template_file_any_unresolved"):
            assert callable(getattr(F, name, None)), f"覆盖层基准别名 {name} 不见了"
        assert F.find_template_file.__wrapped__ is F.find_template_file_unresolved
        assert F.find_template_file_any.__wrapped__ is F.find_template_file_any_unresolved

    def test_literal_codes_is_a_pure_function(self) -> None:
        """Requirement 10.2：不读磁盘、无副作用、同输入恒等输出。"""
        import inspect

        src = inspect.getsource(F._literal_wp_codes_in_filename)
        for forbidden in ("open(", "read_bytes", "read_text", "iterdir", "exists",
                          "glob", "Path("):
            assert forbidden not in src, f"纯函数里出现磁盘访问 {forbidden!r}"
        probe = "A3-7内部往来核对表、A3-8商誉减值测试.xlsx"
        assert F._literal_wp_codes_in_filename(probe) == \
            F._literal_wp_codes_in_filename(probe)

    def test_finder_never_reads_xlsx_bytes(self) -> None:
        """Requirement 10.7：查找器全程不读 xlsx 字节（纯路径解析）。"""
        import inspect

        src = inspect.getsource(F._find_combined_workbook_declaring)
        for forbidden in ("read_bytes", "read_text", "open(", "load_workbook",
                          "ZipFile"):
            assert forbidden not in src, f"查找器读了字节：{forbidden!r}"
        assert "(len(p.name), p.name)" in src, "排序键必须是 (名字长度, 名字)"


class TestProperty10RangeMarkerFailClosed:
    """Property 10: 范围式命名 fail-closed
    （Validates: Requirements 10.3, 10.4）—— **双向**变异证明。"""

    @PBT
    @given(
        marker=st.sampled_from(F._RANGE_MARKERS),
        left=st.sampled_from(["D4-1", "E1-26", "F2-21", "A9-3"]),
        right=st.sampled_from(["D4-4", "E1-32", "F2-26", "A9-7"]),
        tail=st.sampled_from(["", " 营业收入", "存货及跌价准备", "（Leap-常规程序）"]),
    )
    def test_any_filename_with_a_range_marker_yields_the_empty_set(
        self, marker: str, left: str, right: str, tail: str
    ) -> None:
        name = f"{left}{marker}{right}{tail}.xlsx"
        assert F._literal_wp_codes_in_filename(name) == frozenset(), name

    @PBT
    @given(
        codes=st.lists(st.sampled_from(["A3-7", "A3-8", "A2-1", "A2-2", "D4-33"]),
                       min_size=1, max_size=3, unique=True),
        sep=st.sampled_from(["、", "＋", " 与 ", "_"]),
    )
    def test_without_a_range_marker_a_filename_with_codes_is_never_empty(
        self, codes: list[str], sep: str
    ) -> None:
        """变异证明②（反方向）：无范围标记且含至少一个码 ⇒ **恒非空**。

        结构性零必须配非零变异 —— 少了这一条，`_literal_wp_codes_in_filename` 写成
        `return frozenset()` 也会让上一条全绿。
        """
        name = sep.join(codes) + "中文名.xlsx"
        got = F._literal_wp_codes_in_filename(name)
        assert got, f"{name!r} 应非空，实得空集"
        assert set(codes) <= got, (name, sorted(got))

    def test_real_range_books_in_the_repo_are_all_fail_closed(
        self, all_workbook_names: list[str]
    ) -> None:
        """接真目录：现算所有含范围标记的册都返回空集。"""
        ranged = [n for n in all_workbook_names
                  if any(m in n.rsplit(".", 1)[0] for m in F._RANGE_MARKERS)]
        assert ranged, "权威目录里一本范围册都没有 ⇒ 本判据失去分母（口径须重算）"
        for name in ranged:
            assert F._literal_wp_codes_in_filename(name) == frozenset(), name

    def test_a_domain_has_zero_range_style_combined_books(
        self, all_workbook_names: list[str]
    ) -> None:
        """「不实现范围展开」的依据（Requirement 10.9）：A 域现算零本范围式合册。

        这个现算值一旦变成非零，`_RANGE_MARKERS` 的 fail-closed 就会让 A 域范围册
        表现为「不认」—— 那时才需要评估是否实现展开器（届时本判据会打红提醒）。
        """
        a_names = [
            f.name for f in (F.TEMPLATES_DIR / "A").iterdir()
            if f.suffix.lower() in (".xlsx", ".xlsm")
        ]
        ranged_a = [n for n in a_names
                    if any(m in n.rsplit(".", 1)[0] for m in F._RANGE_MARKERS)]
        assert ranged_a == [], (
            f"A 域出现范围式册 {ranged_a} ⇒ 它们会被 fail-closed 判成「不认」，"
            f"请评估是否需要范围展开（本 spec 明确不做，扩展点在 _RANGE_MARKERS）"
        )


class TestProperty11CombinedCriterionRejectsSingleCodeBooks:
    """Property 11: 合册判据拒单码册
    （Validates: Requirements 10.5, 15.4）。"""

    @PBT
    @given(
        code=st.sampled_from(["A3", "A2", "D4", "F2", "A3-7", "B22A"]),
        tail=st.sampled_from(["合并流程程序表", " 收入底稿", "存货", " 审定表"]),
    )
    def test_a_filename_declaring_exactly_one_code_is_never_a_combined_book(
        self, code: str, tail: str
    ) -> None:
        name = f"{code}{tail}.xlsx"
        assert len(F._literal_wp_codes_in_filename(name)) <= 1, name
        assert F._combined_book_covers(name, code) is False, name

    def test_every_single_code_book_in_the_repo_is_rejected(
        self, all_workbook_names: list[str], combined_books: list[str]
    ) -> None:
        """接真目录：现算所有非合册的册都不被 `_combined_book_covers` 认。"""
        singles = [n for n in all_workbook_names if n not in set(combined_books)]
        assert singles, "分母为空"
        for name in singles:
            for code in F._literal_wp_codes_in_filename(name) or {"X1"}:
                assert F._combined_book_covers(name, code) is False, (name, code)

    def test_reverse_mutation_shortcircuiting_the_narrowing_grows_the_diff(self) -> None:
        """反向变异③：短路 `len(declared) < 2` 后，全域解析差集必须**变大**。

        少了这一条，`len(declared) < 2` 写错也看不出来（它会变成恒真装饰）。
        判据用「命中数」而不是跑全域两路解析 —— 后者慢且与 Task 15 的探针重复；
        这里测的是同一个收窄条件的**因果**：短路后被认成合册的文件数必须变多。
        """
        names = [
            f.name for f in (F.TEMPLATES_DIR / "A").iterdir()
            if f.suffix.lower() in (".xlsx", ".xlsm")
        ]
        probe_codes = ["A3-8", "A2-2", "A3", "A2", "A11-2", "A13-2", "A3-2"]

        def covered(fn) -> int:
            return sum(1 for n in names for c in probe_codes if fn(n, c))

        def loosened(filename: str, wp_code: str) -> bool:
            declared = F._literal_wp_codes_in_filename(filename)
            if not declared:            # 只去掉 `< 2` 这一条
                return False
            return wp_code in declared or any(
                wp_code.startswith(code + "-") for code in declared
            )

        strict_hits = covered(F._combined_book_covers)
        loose_hits = covered(loosened)
        assert loose_hits > strict_hits, (
            f"短路收窄条件后命中数未变多（{loose_hits} ≤ {strict_hits}）⇒ "
            f"`len(declared) < 2` 是恒真装饰，不是真收窄"
        )


class TestProperty13AncestorRuleDoesNotCrossCodeBoundaries:
    """Property 13: 祖先规则不越码边界
    （Validates: Requirements 10.6）。"""

    @PBT
    @given(digit=st.sampled_from(["0", "1", "5", "9"]))
    def test_declared_code_plus_a_bare_digit_does_not_match(self, digit: str) -> None:
        """`A3-8` 声明不得命中 `A3-80` —— 那是**另一个码**，不是它的子码。"""
        name = Path(COMBINED_A38).name
        assert F._combined_book_covers(name, f"A3-8{digit}") is False, digit

    @PBT
    @given(suffix=st.sampled_from(["1", "2", "12", "1-3"]))
    def test_declared_code_plus_hyphen_suffix_does_match(self, suffix: str) -> None:
        name = Path(COMBINED_A38).name
        assert F._combined_book_covers(name, f"A3-8-{suffix}") is True, suffix

    def test_boundary_case_a380_resolves_exactly_as_before(self) -> None:
        """越界的负向终态：`A3-80` 不得被合册抢走（它本就无自有载体）。"""
        one, anyf = both_routes("A3-80")
        assert one == "A/A3 合并流程程序表.xlsx", one
        assert anyf is None, anyf


class TestProperty8CombinedDeclaredCodesResolveToTheBookItself:
    """Property 8: 合册声明码解析闭合
    （Validates: Requirements 11.1, 11.2, 12.1）。

    分母是「全部合册 × 其字面声明的全部码」（现算穷举，非抽样 —— 分母很小）。
    """

    def test_every_declared_code_of_every_combined_book_resolves_to_that_book(
        self, combined_books: list[str]
    ) -> None:
        checked = 0
        for name in combined_books:
            # 合册在哪个首字母目录下 —— 由码的首字母决定，两者须一致
            for code in sorted(F._literal_wp_codes_in_filename(name)):
                expected = f"{code[0]}/{name}"
                one, anyf = both_routes(code)
                assert one == expected, (code, one, expected)
                assert anyf == expected, (code, anyf, expected)
                checked += 1
        assert checked >= 4, f"分母过小（{checked}）⇒ 判据形同虚设"


class TestProperty9AncestorCodeResolvesToTheSameBook:
    """Property 9: 祖先码解析到同一本合册
    （Validates: Requirements 10.6, 12.2）。"""

    @PBT
    @given(depth=st.integers(min_value=1, max_value=3))
    def test_descendant_codes_without_their_own_carrier_land_on_the_same_book(
        self, depth: int, combined_books: list[str]
    ) -> None:
        for name in combined_books:
            for code in sorted(F._literal_wp_codes_in_filename(name)):
                child = code + "-1" * depth
                if F.find_template_file_any_unresolved(child) is None:
                    continue            # 该码另有自有载体或本就不可解析
                got_one, got_any = both_routes(child)
                expected = f"{code[0]}/{name}"
                assert got_one == expected, (child, got_one)
                assert got_any == expected, (child, got_any)


class TestMultipleCombinedBooksDeclaringTheSameCode:
    """Requirement 10.7 / 10.8：同一 wp_code 被 >1 本合册声明时**可记录、可复现**。

    该情形现算 **0** 例 ⇒ 判据写成「将来出现时结果不依赖 `iterdir()` 枚举顺序」，
    用构造的临时目录做正面样本。
    """

    def test_currently_zero_codes_are_declared_by_more_than_one_combined_book(
        self, combined_books: list[str]
    ) -> None:
        from collections import Counter

        counter: Counter[str] = Counter()
        for name in combined_books:
            for code in F._literal_wp_codes_in_filename(name):
                counter[code] += 1
        dupes = {c: n for c, n in counter.items() if n > 1}
        assert dupes == {}, (
            f"出现被多本合册声明的码 {dupes} ⇒ 解析取确定性排序首个；"
            f"请人工复核这是模板库重复还是命名冲突（判据下一条已保证结果可复现）"
        )

    def test_ordering_is_deterministic_and_independent_of_iteration_order(
        self, tmp_path: Path, monkeypatch: pytest.MonkeyPatch
    ) -> None:
        """正面样本：构造两本都声明 `Z1-2` 的合册，断言取 `(len(name), name)` 首个。"""
        sub = tmp_path / "Z"
        sub.mkdir()
        # 刻意让「较短的名字」与「字典序较小的名字」分离，证明排序键是二元组
        longer = sub / "Z1-1、Z1-2 aaaaaaaa.xlsx"
        shorter = sub / "Z1-1、Z1-2 bb.xlsx"
        for p in (longer, shorter):
            p.write_bytes(b"x")
        monkeypatch.setattr(F, "TEMPLATES_DIR", tmp_path)
        got = F._find_combined_workbook_declaring("Z1-2")
        assert got == shorter, f"应取名字更短的 {shorter.name}，实得 {got}"

        # 反向：把 `iterdir` 的产出顺序反过来，结果必须不变
        real_iterdir = Path.iterdir

        def reversed_iterdir(self):
            return reversed(list(real_iterdir(self)))

        monkeypatch.setattr(Path, "iterdir", reversed_iterdir)
        assert F._find_combined_workbook_declaring("Z1-2") == shorter, (
            "结果随 iterdir 枚举顺序变化 ⇒ 排序键没起作用"
        )


class TestControlGroupResolutionUnchanged:
    """Requirement 11.7 / 16.4：当前已能解析的码走向**逐个不变**，且**不进**新步骤。"""

    #: 对照组。`A3-7` 在更早步骤命中合册本身；`D2-2` / `E1-5` / `F2-40` 在**范围回退**命中。
    CONTROL = {
        "A3-7": COMBINED_A38,
        "D2-2": "D/D2-1至D2-4  应收账款- 审定表明细表（Leap-常规程序）.xlsx",
        "E1-5": "E/E1-1至E1-11 货币资金- 审定表明细表（Leap-常规程序）.xlsx",
        "F2-40": "F/F2-38至F2-44  存货及跌价准备 -计价测试（Leap-存货程序）.xlsx",
    }

    @pytest.mark.parametrize("code", sorted(CONTROL))
    def test_control_code_resolves_to_its_own_book_on_both_routes(self, code: str) -> None:
        expected = self.CONTROL[code]
        assert both_routes(code) == (expected, expected)

    def test_three_of_them_are_not_a_only_sub_codes(self) -> None:
        """现算实证：`D2-2` / `E1-5` / `F2-40` 的 A-only 正则均 `False`（走通用链）。"""
        for code in ("D2-2", "E1-5", "F2-40"):
            assert F._LEGACY_A_ONLY_SUB_CODE_RE.match(code) is None, code
        assert F._LEGACY_A_ONLY_SUB_CODE_RE.match("A3-7") is not None

    def test_control_codes_never_reach_the_new_step(
        self, monkeypatch: pytest.MonkeyPatch
    ) -> None:
        """把新步骤换成「一进就爆」的探针：对照组一个都不该进来。

        这比「结果相等」更强 —— 结果相等也可能是「进了新步骤但恰好返回同一本册」。
        """
        calls: list[str] = []

        def tripwire(wp_code: str):
            calls.append(wp_code)
            raise AssertionError(f"对照组码 {wp_code!r} 进入了新步骤")

        monkeypatch.setattr(F, "_find_combined_workbook_declaring", tripwire)
        for code, expected in self.CONTROL.items():
            assert both_routes(code) == (expected, expected), code
        assert calls == [], calls

    def test_the_a_only_regex_semantics_is_unchanged(self) -> None:
        """Requirement 11.3：`_LEGACY_A_ONLY_SUB_CODE_RE` 语义不变。

        改成字母类无关会让 `D2-2` / `E1-3` 这类 Excel 子表走子码分支，实测使数百个
        wp_code 变 `None`（D/E/F 十九本范围册零缺陷正靠这条 A-only 正则）。
        """
        assert F._LEGACY_A_ONLY_SUB_CODE_RE.pattern == r"^A\d+-\d+"

    def test_three_shortcuts_are_explicitly_vetoed_in_code_comments(self) -> None:
        """Requirement 11.4 / 11.5：三条捷径写进代码注释显式否决（防下一轮再试）。"""
        src = Path(F.__file__).read_bytes().decode("utf-8")
        seg_start = src.index("# 合册声明码索引（spec workpaper-sync-pure-static-lane")
        seg = src[seg_start:seg_start + 4000]
        assert "_index.json" in seg and "禁" in seg
        assert "wp_code_overrides.json" in seg
        assert "_LEGACY_A_ONLY_SUB_CODE_RE" in seg
        assert "componentType" in seg, "override 表捷径的否决理由须写明（值是 componentType）"

    def test_index_json_route_really_cannot_fix_it(self) -> None:
        """否决①的现算实证：A 子码分支按**文件名**前缀判定，补索引记录无作用。"""
        name = Path(COMBINED_A38).name
        assert F._match_filename_prefix(name, "A3-8") is False
        assert F._match_filename_prefix(name, "A3-7") is True
        index = F._load_index()
        for code in ("A3-8", "A3-8-1", "A2-2"):
            assert [e for e in index if e["wp_code"] == code] == [], code


class TestA38AndA381AreUsableAfterTheFix:
    """Requirement 12：`A3-8` / `A3-8-1` 的期望终态是「**可用**」—— 既解析到册，
    又能指向册内对应 sheet。"""

    def test_both_routes_return_the_combined_book(self) -> None:
        assert both_routes("A3-8") == (COMBINED_A38, COMBINED_A38)

    def test_a381_lands_on_the_same_book_via_the_ancestor_rule(self) -> None:
        """文件名只字面声明到 `A3-8`，`A3-8-1` 只出现在 sheet 名里。"""
        name = Path(COMBINED_A38).name
        declared = F._literal_wp_codes_in_filename(name)
        assert "A3-8-1" not in declared, declared
        assert "A3-8" in declared
        assert both_routes("A3-8-1") == (COMBINED_A38, COMBINED_A38)

    def test_the_book_really_contains_the_two_target_sheets(self) -> None:
        """「可用」的第二半：册内有对应 sheet（sheet 数现算，禁写死）。"""
        from openpyxl import load_workbook

        path = F.TEMPLATES_DIR / COMBINED_A38
        assert path.is_file(), path
        wb = load_workbook(path, read_only=True, data_only=True)
        try:
            sheets = list(wb.sheetnames)
        finally:
            wb.close()
        assert "A3-8商誉减值测试" in sheets, sheets
        assert "A3-8-1可收回金额测试" in sheets, sheets
        assert len(sheets) >= 3, f"该册 sheet 数现算 {len(sheets)}，分母过小须复核"

    def test_before_the_fix_find_template_file_returned_the_wrong_book(self) -> None:
        """前后对照：改动前 `find_template_file("A3-8")` 返回**错册**（静默错答）。

        「改动前」从 git HEAD 版 finder 现取，不是凭记忆写死。
        """
        res = subprocess.run(
            ["git", "show", "HEAD:backend/app/services/wp_template_finder.py"],
            cwd=_BACKEND.parent, capture_output=True, check=False,
        )
        if res.returncode != 0 or not res.stdout:
            pytest.skip("git show 取不到 HEAD 版 finder")
        import types

        head = types.ModuleType("_wp_template_finder_head")
        head.__file__ = str(Path(F.__file__))
        sys.modules[head.__name__] = head
        exec(compile(res.stdout.decode("utf-8"), head.__file__, "exec"), head.__dict__)

        assert head.find_template_file_any_unresolved("A3-8") is None, (
            "改动前 `find_template_file_any('A3-8')` 应为 None"
        )
        wrong = head.find_template_file_unresolved("A3-8")
        assert wrong is not None
        assert Path(wrong).name == "A3 合并流程程序表.xlsx", (
            f"改动前应返回父程序表这个错册，实得 {Path(wrong).name!r}"
        )
        # 改动后：两路都是合册
        assert both_routes("A3-8") == (COMBINED_A38, COMBINED_A38)

    def test_two_defect_denominators_are_counted_separately(self) -> None:
        """Requirement 12.5：两个分母**分开**记账，合并成单一「3」= 不合格。

        按文件名扫描永远只得 ②（现算 2 个）；把 `A3-8-1` 混进去会让扫描器判据写错、
        被误判成漏抓。
        """
        declared = F._literal_wp_codes_in_filename(Path(COMBINED_A38).name) | \
            F._literal_wp_codes_in_filename(Path(COMBINED_A22).name)
        fixed = {"A3-8", "A3-8-1", "A2-2"}
        by_filename = sorted(fixed & declared)
        by_sheet_only = sorted(fixed - declared)
        assert by_filename == ["A2-2", "A3-8"], by_filename
        assert by_sheet_only == ["A3-8-1"], by_sheet_only
        assert len(by_filename) + len(by_sheet_only) == len(fixed)

    def test_second_step_is_registered_as_a_successor_not_done_here(self) -> None:
        """Requirement 12.6：只做第一步（解析层）；第二步（宿主传中文字面 sheet 名）
        登记为后继。顺序不可颠倒的实证 = 改动前 `_any("A3-8")` 为 `None`
        ⇒ 第二步单独做完零收益（上一条测试已现算该 `None`）。"""
        slice_path = (_BACKEND / "data"
                      / "workpaper_sync_abcs_cycle_manifest_slice.json")
        doc = json.loads(slice_path.read_bytes().decode("utf-8"))
        entry = next(
            e for e in doc["independent_entries"]
            if e["entry_id"] == "xlsx/gt-a38-goodwill-impairment"
        )
        ref = entry["template_ref"]
        # 第二步的两个字段仍是宿主层的**声明**（本 spec 不改它们的语义）
        assert ref["resolution_kind"] == "literal_sheet_name"
        assert ref["sheet_name_literal"] == "A3-8商誉减值测试"


class TestA22IsBookResolvableOnly:
    """Requirement 13：`A2-2` 的终态是「**仅册可解析**」，不承诺可用。"""

    def test_both_routes_return_the_declaring_book(self) -> None:
        assert both_routes("A2-2") == (COMBINED_A22, COMBINED_A22)

    def test_sheet_level_supply_is_recorded_as_a_separate_gap(self) -> None:
        """如实记录 sheet 层事实：两数**现算、禁写死**；结论是后者为 0 ⇒ 仍无供给。"""
        from openpyxl import load_workbook

        path = F.TEMPLATES_DIR / COMBINED_A22
        wb = load_workbook(path, read_only=True, data_only=True)
        try:
            sheets = list(wb.sheetnames)
        finally:
            wb.close()
        starting_with_a22 = [s for s in sheets if s.startswith("A2-2")]
        assert len(sheets) > 0
        assert starting_with_a22 == [], (
            f"该册出现以 `A2-2` 起头的 sheet {starting_with_a22} ⇒ sheet 层供给缺口已补，"
            f"`A2-2` 可从「仅册可解析」升为「可用」，请更新本判据与登记"
        )

    def test_a22_is_explicitly_not_claimed_usable(self) -> None:
        """判据显式区分「册可解析」（承诺）与「可用」（不承诺）。

        把 `A2-2` 与 `A3-8` / `A3-8-1` 合写成「修好就能用」= 制造假绿。
        """
        assert both_routes("A2-2")[0] is not None        # 册可解析：承诺
        from openpyxl import load_workbook

        wb = load_workbook(F.TEMPLATES_DIR / COMBINED_A22, read_only=True,
                           data_only=True)
        try:
            assert not any(s.startswith("A2-2") for s in wb.sheetnames)
        finally:
            wb.close()


class TestAstOrderInvariantIsNotBroken:
    """B-7 AST 顺序不变式 —— **EXAMPLE 不是属性**（单一结构断言，既有守卫已覆盖）。

    Requirement 16.1 / 16.2：两处插入**不得**移动 `_resolve_most_specific_docx` 与
    `_LEGACY_A_ONLY_SUB_CODE_RE.match` 两个调用的位置与先后关系。
    行号由 AST **现算**，禁写死。
    """

    def test_docx_probe_still_precedes_the_a_only_regex_inside_find_any(self) -> None:
        import ast

        src = Path(F.__file__).read_bytes().decode("utf-8")
        tree = ast.parse(src)
        fn = next(
            n for n in ast.walk(tree)
            if isinstance(n, ast.FunctionDef) and n.name == "find_template_file_any"
        )
        docx_line = a_only_line = None
        for node in ast.walk(fn):
            if isinstance(node, ast.Call):
                target = node.func
                if isinstance(target, ast.Name) and \
                        target.id == "_resolve_most_specific_docx":
                    docx_line = node.lineno if docx_line is None else docx_line
                if isinstance(target, ast.Attribute) and target.attr == "match" and \
                        isinstance(target.value, ast.Name) and \
                        target.value.id == "_LEGACY_A_ONLY_SUB_CODE_RE":
                    a_only_line = node.lineno if a_only_line is None else a_only_line
        assert docx_line is not None, "`_resolve_most_specific_docx` 调用不见了"
        assert a_only_line is not None, "`_LEGACY_A_ONLY_SUB_CODE_RE.match` 调用不见了"
        assert docx_line < a_only_line, (docx_line, a_only_line)

    def test_the_existing_ast_guard_test_still_passes(self) -> None:
        """跑既有守卫本体（它是这条不变式的权威判据，本文件只保证不被破坏）。"""
        res = subprocess.run(
            [sys.executable, "-m", "pytest", "-q", "--no-header", "-p", "no:cacheprovider",
             "-W", "ignore",
             # 🔴 nodeid 含类名：该判据是 `TestProperty40MostSpecificSubCode` 的方法，
             #    不是顶层函数（少了类名 pytest 报 `not found` 而**不是**失败 ⇒
             #    会被误读成「守卫不存在」）。
             "tests/workpaper_sync/test_task58_word_canonical_resolver.py"
             "::TestProperty40MostSpecificSubCode"
             "::test_finder_no_longer_uses_a_only_regex_for_sub_code_decision"],
            cwd=_BACKEND, capture_output=True, check=False,
        )
        out = (res.stdout + res.stderr).decode("utf-8", "replace")
        assert res.returncode == 0, out[-3000:]

    def test_new_step_is_inserted_after_the_two_prefix_attempts(self) -> None:
        """插入位置：`find_template_file_any` 里在两次同名前缀尝试**之后**、
        `return None` **之前**。"""
        import inspect

        src = inspect.getsource(F.find_template_file_any_unresolved)
        i_index_try = src.index("xlsx_matches = [")
        i_disk_try = src.index("for f in sorted(subdir.iterdir()):")
        i_new = src.index("_find_combined_workbook_declaring(wp_code)")
        i_return_none = src.rindex("return None")
        assert i_index_try < i_disk_try < i_new < i_return_none, (
            i_index_try, i_disk_try, i_new, i_return_none
        )

    def test_new_step_sits_between_range_fallback_and_last_resort(self) -> None:
        """插入位置：`find_template_file` 里在「{主码}-N至{主码}-M」范围回退**之后**、
        终极回退 `startswith(主码 + " ")` **之前**。"""
        import inspect

        src = inspect.getsource(F.find_template_file_unresolved)
        i_range = src.index('and "至" in f.name')
        i_new = src.index("_find_combined_workbook_declaring(wp_code)")
        i_last = src.index('f.name.startswith(primary + " ")')
        assert i_range < i_new < i_last, (i_range, i_new, i_last)


class TestCombinedBookGlobalFactsUnchanged:
    """Requirement 16.3：多码册全域事实现算并保持不变（全部现算、禁写死）。"""

    def test_totals_forms_and_no_overlap(self, all_workbook_names: list[str]) -> None:
        import re as _re
        from collections import Counter

        code_re = F._FILENAME_CODE_RE
        multi: list[tuple[str, str]] = []
        for sub in sorted(p for p in F.TEMPLATES_DIR.iterdir() if p.is_dir()):
            for f in sorted(sub.iterdir()):
                if f.suffix.lower() not in (".xlsx", ".xlsm"):
                    continue
                stem = f.name.rsplit(".", 1)[0]
                # 「多码册」口径：**忽略**范围标记后仍有 ≥2 个字面码
                if len(set(code_re.findall(stem))) >= 2:
                    multi.append((sub.name, f.name))
        by_letter = Counter(letter for letter, _n in multi)
        ranged = [n for _l, n in multi
                  if any(m in n.rsplit(".", 1)[0] for m in F._RANGE_MARKERS)]
        dunhao = [n for _l, n in multi
                  if "、" in n.rsplit(".", 1)[0] and n not in set(ranged)]
        other = [n for _l, n in multi
                 if n not in set(ranged) and n not in set(dunhao)]

        assert len(ranged) + len(dunhao) + len(other) == len(multi), (
            "分形态之和 ≠ 总数 ⇒ 口径有重叠或遗漏"
        )
        assert set(ranged) & set(dunhao) == set(), "顿号集合与范围集合有交叠"
        assert other == [], f"出现第三种分隔形态 {other} ⇒ 口径须重算"
        # 分形态的域分布：顿号全在 A、范围全在 D/E/F（现算结论，变了就打红）
        assert all(n.startswith("A") for n in dunhao), dunhao
        assert all(n[0] in "DEF" for n in ranged), sorted(ranged)[:5]
        assert set(by_letter) == {"A", "D", "E", "F"}, dict(by_letter)
        assert len(dunhao) == by_letter["A"], (len(dunhao), by_letter["A"])
        assert _re.match(r"^[A-Z]", multi[0][1])

    def test_def_range_books_have_zero_resolution_defects(self) -> None:
        """「不改 A-only 正则」的依据：D/E/F 范围册的码现算零缺陷。"""
        bad: list[str] = []
        for sub in ("D", "E", "F"):
            for f in sorted((F.TEMPLATES_DIR / sub).iterdir()):
                if f.suffix.lower() not in (".xlsx", ".xlsm"):
                    continue
                stem = f.name.rsplit(".", 1)[0]
                if not any(m in stem for m in F._RANGE_MARKERS):
                    continue
                for code in F._FILENAME_CODE_RE.findall(stem):
                    one, anyf = both_routes(code)
                    if one is None or anyf is None:
                        bad.append(f"{code} → ({one}, {anyf})")
        assert bad == [], f"范围册的码出现解析不到 {bad[:8]}"


# ═══════════════════════════════════════════════════════════════════════════
# §8 全域前后差集（Requirement 15 / Property 12）
# ═══════════════════════════════════════════════════════════════════════════
#
# 🔴 **量化域就是全域分母本身 ⇒ 用穷举而非 `hypothesis` 抽样**（穷举严格强于
#    `max_examples=5`；降级成随机抽样反而弱化判据）。
# 🔴 「改动前」由 **git HEAD 版 finder** 现算，不依赖任何一次性探针的快照文件 ——
#    判据因此是**自包含**的，探针删掉后它仍然可跑。

#: 预期被修复的三码及其预期新值（相对 `backend/wp_templates/`）。
EXPECTED_NEW = {
    "A3-8": COMBINED_A38,
    "A3-8-1": COMBINED_A38,
    "A2-2": COMBINED_A22,
}


def wp_code_universe() -> list[str]:
    """全域 wp_code 分母 = 三源之并 ∪ 显式追加项（现算，禁写死分母数）。

    口径（Task 1.1 一次性钉死）：
      1. `wp_templates/_index.json` 的 `wp_code` 去重（剔 `_ref` 占位）；
      2. `app/data/wp_code_overrides.json` 的全部键；
      3. `wp_render_schema/*.yaml` 的 stem —— 现算存在**两处**同名目录
         （`app/data/wp_render_schema` 与 `data/ledger_adapters/wp_render_schema`），
         裁决为**两处都收且递归**。理由：本分母唯一用途是零回归的全域差集，
         分母越大判据越强，漏收才会让被顺带改走向的码藏身 ⇒ 取上界。
      4. 🔴 **显式追加项**：`A2-2` 现算**不在**三源之并里 —— 不追加它，这个码的修复
         就根本不在量化域内。
    """
    codes: set[str] = set()
    idx = json.loads((_BACKEND / "wp_templates" / "_index.json")
                     .read_bytes().decode("utf-8"))
    for e in idx.get("files") or []:
        code = str(e.get("wp_code") or "").strip()
        if code and code != "_ref":
            codes.add(code)
    ov = json.loads((_BACKEND / "app" / "data" / "wp_code_overrides.json")
                    .read_bytes().decode("utf-8"))
    codes |= {str(k).strip() for k in ov if str(k).strip()}
    for rel in (_BACKEND / "app" / "data" / "wp_render_schema",
                _BACKEND / "data" / "ledger_adapters" / "wp_render_schema"):
        if rel.exists():
            codes |= {p.stem for p in rel.rglob("*.yaml")}
    codes |= set(EXPECTED_NEW)
    assert len(codes) > 1000, f"全域分母只有 {len(codes)} ⇒ 口径可疑"
    return sorted(codes)


@pytest.fixture(scope="module")
def head_finder() -> Any:
    """git HEAD 版 finder（真实「改造前」）。"""
    import types

    res = subprocess.run(
        ["git", "show", "HEAD:backend/app/services/wp_template_finder.py"],
        cwd=_BACKEND.parent, capture_output=True, check=False,
    )
    if res.returncode != 0 or not res.stdout:
        pytest.skip("git show 取不到 HEAD 版 finder")
    mod = types.ModuleType("_wp_finder_head_global_diff")
    mod.__file__ = str(Path(F.__file__))
    sys.modules[mod.__name__] = mod
    exec(compile(res.stdout.decode("utf-8"), mod.__file__, "exec"), mod.__dict__)
    return mod


def resolve_universe(finder: Any, codes: list[str]) -> dict[str, tuple]:
    tdir = finder.TEMPLATES_DIR.resolve()

    def r(p):
        if p is None:
            return None
        try:
            return str(Path(p).resolve().relative_to(tdir)).replace("\\", "/")
        except Exception:  # noqa: BLE001
            return str(p).replace("\\", "/")

    out: dict[str, tuple] = {}
    for code in codes:
        try:
            one = r(finder.find_template_file_unresolved(code))
        except Exception as exc:  # noqa: BLE001
            one = f"!!{type(exc).__name__}"
        try:
            anyf = r(finder.find_template_file_any_unresolved(code))
        except Exception as exc:  # noqa: BLE001
            anyf = f"!!{type(exc).__name__}"
        out[code] = (one, anyf)
    return out


class TestProperty12GlobalResolutionDiffIsExactlyTheThreeCodes:
    """Property 12: 全域解析差集恰为预期三码
    （Validates: Requirements 15.1, 15.3）。

    🔴 零回归判据**必须**是全域差集；以抽样若干码替代 = 不合格。
    """

    def test_diff_is_exactly_the_expected_three_codes(self, head_finder: Any) -> None:
        codes = wp_code_universe()
        before = resolve_universe(head_finder, codes)
        after = resolve_universe(F, codes)
        diff = {c: (before[c], after[c]) for c in codes if before[c] != after[c]}
        assert sorted(diff) == sorted(EXPECTED_NEW), (
            f"全域分母 {len(codes)} 个码；差集应恰为 {sorted(EXPECTED_NEW)}，"
            f"实得 {sorted(diff)}；"
            f"多出 {sorted(set(diff) - set(EXPECTED_NEW))} 缺少 "
            f"{sorted(set(EXPECTED_NEW) - set(diff))}"
        )
        for code, expected in EXPECTED_NEW.items():
            assert after[code] == (expected, expected), (code, after[code])
            # 改动前：一路错册 / 一路 None（两种失效形态都在）
            assert before[code][1] is None, (code, before[code])
            assert before[code][0] != expected, (code, before[code])

    def test_reverse_mutation_over_the_whole_universe_grows_the_diff(
        self, head_finder: Any
    ) -> None:
        """反向变异③（全域口径）：短路 `len(declared) < 2` 后差集必须**变大**。

        少了它，`len(declared) < 2` 写错也看不出来（它会变成恒真装饰）。
        """
        codes = wp_code_universe()
        before = resolve_universe(head_finder, codes)
        strict = resolve_universe(F, codes)
        strict_diff = {c for c in codes if before[c] != strict[c]}

        original = F._combined_book_covers

        def loosened(filename: str, wp_code: str) -> bool:
            declared = F._literal_wp_codes_in_filename(filename)
            if not declared:            # 只去掉 `< 2` 这一条，其余判据不动
                return False
            return wp_code in declared or any(
                wp_code.startswith(code + "-") for code in declared
            )

        F._combined_book_covers = loosened          # type: ignore[assignment]
        try:
            loose = resolve_universe(F, codes)
        finally:
            F._combined_book_covers = original      # type: ignore[assignment]
        loose_diff = {c for c in codes if before[c] != loose[c]}

        assert len(loose_diff) > len(strict_diff), (
            f"短路收窄条件后全域差集未变大（{len(loose_diff)} ≤ {len(strict_diff)}）"
            f"⇒ `len(declared) < 2` 是恒真装饰"
        )
        assert strict_diff < loose_diff, "严格差集应是宽松差集的真子集"

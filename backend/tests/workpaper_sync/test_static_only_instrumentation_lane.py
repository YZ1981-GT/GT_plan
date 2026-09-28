"""纯静态 instrumentation 通道（Lane A）的判据。

spec: workpaper-sync-pure-static-lane-and-combined-workbook-resolution

═══ 本文件守什么 ═══

平台 instrumentation 管线原先**六处**预设「每个 entry 至少有一张动态行表」。Lane A 的
总策略是**旁路而非放宽**：3 处旁路（新增静态类型 / 静态注入器 / 静态 payload 构建器）
+ 3 处加分派臂（substrate 第三臂 / 身份 binding 静态臂 / 观测清册静态形态），
**零处**放宽既有校验。

⇒ 本文件的每个断言都成对出现：「静态通道成立」+「既有动态通道未被放宽」。
后者是对照组（正面判据），缺了它「零回归」就只是一句口号。

判据纪律：一切计数现算（`REQUIRED_GT_SYNC_KEYS` 等按常量取），禁写死键数与行号。
"""
from __future__ import annotations

import dataclasses
import io
import json
import os
import re
import sys
import zipfile
from pathlib import Path
from typing import Any

import pytest

_BACKEND = Path(__file__).resolve().parents[2]
if str(_BACKEND) not in sys.path:
    sys.path.insert(0, str(_BACKEND))
os.environ.setdefault("DB_DISABLE_SSL", "True")

from app.services.workpaper_sync import excel_instrumentation as _xi  # noqa: E402
from app.services.workpaper_sync.excel_instrumentation import (  # noqa: E402
    ExcelInstrumentationSpec,
    ExcelStaticOnlyInstrumentationSpec,
    InstrumentationError,
    StaticRegionSpec,
)

#: 纯静态 spec **禁**出现的行表几何字段（Requirement 1.1）。
ROW_GEOMETRY_FIELDS = (
    "first_data_row",
    "last_data_row",
    "footer_row",
    "uuid_col",
    "managed_last_col",
    "table_name",
)


def region(**overrides: Any) -> StaticRegionSpec:
    """合法静态区的最小样例；`overrides` 用于逐项构造非法输入。"""
    base: dict[str, Any] = {
        "sheet_key": "a511-audit",
        "excel_name": "A5-1-1 现金流量表审定表",
        "template_id": "A511",
        "defined_name": "GT_MANAGED_REGION_A511",
        "managed_ref": "$D$8:$H$15",
        "table_key": "a51_audit_table",
    }
    base.update(overrides)
    return StaticRegionSpec(**base)


def spec_of(*regions: StaticRegionSpec) -> ExcelStaticOnlyInstrumentationSpec:
    return ExcelStaticOnlyInstrumentationSpec(
        entry_id="xlsx/gt-a51-cashflow-audit",
        template_id="A511",
        template_relative_path="A/A5-1 现金流量表审计.xlsx",
        static_regions=regions,
    )


# ═══════════════════════════════════════════════════════════════════════════
# §1 Static_Spec 的校验面（Requirements 1.1 ~ 1.10）
# ═══════════════════════════════════════════════════════════════════════════


class TestStaticSpecShape:
    """字段集合与继承关系 —— 「零行表几何」是结构事实而不是约定。"""

    def test_static_region_spec_has_exactly_six_fields(self) -> None:
        names = [f.name for f in dataclasses.fields(StaticRegionSpec)]
        assert set(names) == {
            "sheet_key", "excel_name", "template_id",
            "defined_name", "managed_ref", "table_key",
        }, names
        assert len(names) == 6, names

    def test_neither_static_type_carries_row_geometry(self) -> None:
        for cls in (StaticRegionSpec, ExcelStaticOnlyInstrumentationSpec):
            fields = {f.name for f in dataclasses.fields(cls)}
            leaked = fields & set(ROW_GEOMETRY_FIELDS)
            assert not leaked, f"{cls.__name__} 泄漏了行表几何字段: {sorted(leaked)}"

    def test_static_spec_is_sibling_not_subclass(self) -> None:
        """兄弟关系：静态类型不得继承既有类（继承会把行几何校验带进来）。"""
        assert not issubclass(ExcelStaticOnlyInstrumentationSpec, ExcelInstrumentationSpec)
        assert not issubclass(ExcelInstrumentationSpec, ExcelStaticOnlyInstrumentationSpec)
        assert ExcelStaticOnlyInstrumentationSpec.__mro__[1:] == (object,)

    def test_both_static_types_are_frozen(self) -> None:
        for cls in (StaticRegionSpec, ExcelStaticOnlyInstrumentationSpec):
            assert dataclasses.fields(cls)  # 是 dataclass
            with pytest.raises(dataclasses.FrozenInstanceError):
                obj = region() if cls is StaticRegionSpec else spec_of(region())
                object.__setattr__  # noqa: B018  — 明确用 setattr 而非绕过
                setattr(obj, dataclasses.fields(cls)[0].name, "x")

    def test_excel_name_is_documented_as_build_time_selector(self) -> None:
        """`excel_name` 的语义是构建期选择器而非运行时锚点（Requirement 1.3）。"""
        import inspect

        src = inspect.getsource(StaticRegionSpec)
        assert "构建期" in src and "运行时锚点" in src, src


class TestStaticSpecRejections:
    """五类拒收各一例；异常消息须含字段名与实得值（Requirements 1.4 ~ 1.7）。"""

    def test_rejects_blank_entry_id(self) -> None:
        with pytest.raises(InstrumentationError) as exc:
            ExcelStaticOnlyInstrumentationSpec(
                entry_id="   ", template_id="A511",
                template_relative_path="A/x.xlsx", static_regions=(region(),),
            )
        assert "entry_id" in str(exc.value)

    def test_rejects_empty_static_regions(self) -> None:
        with pytest.raises(InstrumentationError) as exc:
            ExcelStaticOnlyInstrumentationSpec(
                entry_id="e", template_id="A511",
                template_relative_path="A/x.xlsx", static_regions=(),
            )
        assert "static_regions" in str(exc.value)

    @pytest.mark.parametrize("bad_name", ["1bad", "has space", "有中文", "", "a-b"])
    def test_rejects_illegal_defined_name_charset(self, bad_name: str) -> None:
        with pytest.raises(InstrumentationError) as exc:
            spec_of(region(defined_name=bad_name))
        msg = str(exc.value)
        assert "defined_name" in msg and repr(bad_name) in msg, msg

    @pytest.mark.parametrize("ref_like", ["A1", "AB12", "R1C1", "r10c3"])
    def test_rejects_defined_name_that_looks_like_a1_reference(self, ref_like: str) -> None:
        with pytest.raises(InstrumentationError) as exc:
            spec_of(region(defined_name=ref_like))
        msg = str(exc.value)
        assert "defined_name" in msg and repr(ref_like) in msg, msg

    @pytest.mark.parametrize("bad_ref", ["D8:H15", "$D8:$H$15", "$D$8", "D8", "$AAAA$1:$B$2"])
    def test_rejects_non_absolute_rectangle(self, bad_ref: str) -> None:
        with pytest.raises(InstrumentationError) as exc:
            spec_of(region(managed_ref=bad_ref))
        msg = str(exc.value)
        assert "managed_ref" in msg and repr(bad_ref) in msg, msg

    @pytest.mark.parametrize("bad_ref", ["$H$15:$D$8", "$D$15:$H$8", "$H$8:$D$15"])
    def test_rejects_rectangle_whose_bottom_right_is_not_ge_top_left(self, bad_ref: str) -> None:
        with pytest.raises(InstrumentationError) as exc:
            spec_of(region(managed_ref=bad_ref))
        msg = str(exc.value)
        assert "managed_ref" in msg and repr(bad_ref) in msg, msg

    def test_accepts_degenerate_single_cell_rectangle(self) -> None:
        """右下 == 左上 是合法的（判据是 ≥ 而不是 >）。"""
        assert spec_of(region(managed_ref="$D$8:$D$8")).static_regions


class TestStaticSpecThreeKeyUniqueness:
    """三键重复的三种组合各一例（Requirement 1.8）。"""

    def test_duplicate_sheet_key_is_rejected_at_construction(self) -> None:
        with pytest.raises(InstrumentationError) as exc:
            spec_of(region(), region(defined_name="GT_B", table_key="tk2"))
        assert "sheet_key" in str(exc.value)

    def test_duplicate_defined_name_is_rejected_at_construction(self) -> None:
        with pytest.raises(InstrumentationError) as exc:
            spec_of(region(), region(sheet_key="k2", table_key="tk2"))
        assert "defined_name" in str(exc.value)

    def test_duplicate_table_key_is_rejected_at_construction(self) -> None:
        with pytest.raises(InstrumentationError) as exc:
            spec_of(region(), region(sheet_key="k2", defined_name="GT_B"))
        assert "table_key" in str(exc.value)

    def test_three_distinct_regions_are_accepted(self) -> None:
        got = spec_of(
            region(),
            region(sheet_key="k2", defined_name="GT_B", table_key="tk2"),
            region(sheet_key="k3", defined_name="GT_C", table_key="tk3"),
        )
        assert len(got.static_regions) == 3


class TestRowJudgementsAreOutsideTheStaticValidationSurface:
    """负向：行区间 / footer 行 / UUID 列位置三类判据**不在**静态校验面。

    这是 Requirement 1.9 的可执行形式 —— 只给静态字段即通过，说明静态类型没有偷偷
    继承行几何判据；反过来也证明「静态校验面被限定住了」。
    """

    def test_construction_succeeds_with_static_fields_only(self) -> None:
        got = spec_of(region())
        assert got.semantic_version == "1.0.0"
        assert not hasattr(got, "managed_range")
        assert not hasattr(got, "table_ref")
        assert not hasattr(got, "row_count")

    def test_post_init_source_mentions_none_of_the_row_judgements(self) -> None:
        import inspect

        src = inspect.getsource(ExcelStaticOnlyInstrumentationSpec.__post_init__)
        for token in ("footer_row", "uuid_col", "managed_last_col",
                      "first_data_row", "last_data_row", "table_name"):
            assert token not in src, f"静态校验面泄漏了行几何判据 {token!r}"

    def test_static_spec_has_no_row_uuid_generator(self) -> None:
        assert not hasattr(ExcelStaticOnlyInstrumentationSpec, "row_uuid")
        assert not hasattr(StaticRegionSpec, "row_uuid")


class TestExistingDynamicSpecWasNotRelaxed:
    """对照组：`ExcelInstrumentationSpec` 的字段集合与校验项逐条保持改动前形态。

    Requirement 1.2 / 7.1 的可执行形式。少了这组，「零处放宽」只是一句口号。
    """

    def test_dynamic_spec_still_requires_every_row_geometry_field(self) -> None:
        required = {
            f.name for f in dataclasses.fields(ExcelInstrumentationSpec)
            if f.default is dataclasses.MISSING
            and f.default_factory is dataclasses.MISSING  # type: ignore[misc]
        }
        missing = set(ROW_GEOMETRY_FIELDS) - required
        assert not missing, f"行表几何被改成可选: {sorted(missing)}"

    @pytest.mark.parametrize(
        "kwargs,needle",
        [
            ({"table_name": "has space"}, "displayName"),
            ({"first_data_row": 0}, "受管行区间"),
            ({"last_data_row": 3}, "受管行区间"),
            ({"footer_row": 5}, "footer_row"),
            ({"uuid_col": "A"}, "UUID"),
        ],
    )
    def test_dynamic_spec_post_init_rejections_all_still_fire(
        self, kwargs: dict[str, Any], needle: str
    ) -> None:
        base: dict[str, Any] = {
            "entry_id": "e", "template_id": "T", "template_relative_path": "A/x.xlsx",
            "managed_sheet": "S", "first_data_row": 5, "last_data_row": 20,
            "footer_row": 21, "managed_last_col": "F", "uuid_col": "H",
            "table_name": "GtRow",
        }
        assert ExcelInstrumentationSpec(**base)          # 基线可构造
        base.update(kwargs)
        with pytest.raises(InstrumentationError) as exc:
            ExcelInstrumentationSpec(**base)
        assert needle in str(exc.value), str(exc.value)

    def test_dynamic_spec_still_keeps_parasitic_static_sheets_field(self) -> None:
        """寄生字段仍在 —— 既有静态区全部挂在动态 primary 上，不得顺手删。"""
        names = {f.name for f in dataclasses.fields(ExcelInstrumentationSpec)}
        assert "static_sheets" in names
        assert "transposed_sheets" in names


# ═══════════════════════════════════════════════════════════════════════════
# §2 Static_Injector（Requirements 2.1 ~ 2.14 / Properties 1 ~ 3）
# ═══════════════════════════════════════════════════════════════════════════

from hypothesis import HealthCheck, given, settings  # noqa: E402
from hypothesis import strategies as st  # noqa: E402

#: 项目铁律：`max_examples=5`，禁 hypothesis 默认 100。
PBT = settings(
    max_examples=5,
    deadline=None,
    suppress_health_check=[HealthCheck.function_scoped_fixture, HealthCheck.too_slow],
)

#: sheet 名池刻意含空格/连字符/中文括号 —— `_quote_sheet_name` 只在命中这些字符时加引号，
#: 只用纯字母名会让「ref 带引号」这条路径从不被执行。
SHEET_NAME_POOL = ("Alpha", "Beta 2", "Gamma-3", "德尔塔（四）", "Eps")


def build_source_workbook(sheet_names: list[str]) -> bytes:
    """合法 xlsx 源册：每张 sheet 都有值与公式，用于逐格等价比对。"""
    from openpyxl import Workbook

    wb = Workbook()
    wb.remove(wb.active)
    for si, name in enumerate(sheet_names, start=1):
        ws = wb.create_sheet(title=name)
        for row in range(1, 9):
            ws.cell(row=row, column=1, value=f"{name}-label-{row}")
            ws.cell(row=row, column=2, value=row * si)
            ws.cell(row=row, column=3, value=f"=B{row}*2")
    buf = io.BytesIO()
    wb.save(buf)
    return buf.getvalue()


def cells_by_sheet(data: bytes) -> dict[str, dict[str, Any]]:
    """每张非 `_GT_SYNC` sheet 的逐格值（公式以字符串形式读出）。"""
    from openpyxl import load_workbook

    wb = load_workbook(io.BytesIO(data), data_only=False)
    out: dict[str, dict[str, Any]] = {}
    for ws in wb.worksheets:
        if ws.title == _xi.GT_SYNC_SHEET_NAME:
            continue
        out[ws.title] = {
            cell.coordinate: cell.value
            for row in ws.iter_rows()
            for cell in row
            if cell.value is not None
        }
    return out


def columns_by_sheet(data: bytes) -> dict[str, list[int]]:
    """每张非 `_GT_SYNC` sheet 出现过的列号集合（用于「零新增隐藏列」判据）。"""
    from openpyxl import load_workbook

    wb = load_workbook(io.BytesIO(data), data_only=False)
    return {
        ws.title: sorted({c.column for row in ws.iter_rows() for c in row if c.value is not None})
        for ws in wb.worksheets
        if ws.title != _xi.GT_SYNC_SHEET_NAME
    }


def table_part_count(data: bytes) -> int:
    with zipfile.ZipFile(io.BytesIO(data)) as zf:
        return len([n for n in zf.namelist() if n.startswith("xl/tables/")])


@st.composite
def source_and_spec(draw) -> tuple[bytes, ExcelStaticOnlyInstrumentationSpec]:
    """随机合法源册 × 随机合法静态 spec 集合（sheet 数 / 区数 / 矩形位置全随机）。"""
    sheet_count = draw(st.integers(min_value=1, max_value=len(SHEET_NAME_POOL)))
    names = list(SHEET_NAME_POOL[:sheet_count])
    region_count = draw(st.integers(min_value=1, max_value=sheet_count))
    regions: list[StaticRegionSpec] = []
    for i in range(region_count):
        c1 = draw(st.integers(min_value=1, max_value=5))
        c2 = draw(st.integers(min_value=c1, max_value=c1 + 4))
        r1 = draw(st.integers(min_value=1, max_value=6))
        r2 = draw(st.integers(min_value=r1, max_value=r1 + 6))
        from openpyxl.utils import get_column_letter

        regions.append(StaticRegionSpec(
            sheet_key=f"key-{i}",
            excel_name=names[i],
            template_id=f"T{i}",
            defined_name=f"GT_MANAGED_REGION_T{i}",
            managed_ref=f"${get_column_letter(c1)}${r1}:${get_column_letter(c2)}${r2}",
            table_key=f"tk-{i}",
        ))
    spec = ExcelStaticOnlyInstrumentationSpec(
        entry_id="xlsx/gt-pbt-static",
        template_id="T0",
        template_relative_path="A/pbt.xlsx",
        static_regions=tuple(regions),
    )
    return build_source_workbook(names), spec


@pytest.fixture(scope="module")
def gate() -> Any:
    return _xi.ExcelIdentityCarrierGate.load()


#: A5-1 现算的载体清单（2 项）—— 刻意**不**用 `gate.allowed_carriers`（现算 4 项），
#: 按 gate 推导即语义过度声明（design §1.7）。
STATIC_CARRIERS = ("defined_name", "hidden_sheet")


def inject(source: bytes, spec: ExcelStaticOnlyInstrumentationSpec, gate: Any):
    return _xi.instrument_workbook_bytes_static_only(
        source, spec, gate=gate, identity_carriers=STATIC_CARRIERS
    )


class TestProperty1VisibleSurfaceUnchanged:
    """Property 1: 静态注入不改可见业务面（Validates: Requirements 2.12）。"""

    @PBT
    @given(payload=source_and_spec())
    def test_non_gt_sync_cells_are_equal_cell_by_cell(self, payload, gate) -> None:
        source, spec = payload
        got = inject(source, spec, gate)
        assert cells_by_sheet(got.instrumented_bytes) == cells_by_sheet(source)

    def test_sheet_parts_are_byte_identical(self, gate) -> None:
        """更强的同源判据：本注入器**一个 sheet 部件的字节都不动**。"""
        source, spec = build_source_workbook(list(SHEET_NAME_POOL)), None
        from openpyxl.utils import get_column_letter

        spec = ExcelStaticOnlyInstrumentationSpec(
            entry_id="e", template_id="T0", template_relative_path="A/x.xlsx",
            static_regions=tuple(
                StaticRegionSpec(
                    sheet_key=f"k{i}", excel_name=n, template_id=f"T{i}",
                    defined_name=f"GT_R_T{i}",
                    managed_ref=f"$A$1:${get_column_letter(3)}$8", table_key=f"t{i}",
                )
                for i, n in enumerate(SHEET_NAME_POOL)
            ),
        )
        got = inject(source, spec, gate)
        with zipfile.ZipFile(io.BytesIO(source)) as z0, \
                zipfile.ZipFile(io.BytesIO(got.instrumented_bytes)) as z1:
            changed = [n for n in z0.namelist() if z0.read(n) != z1.read(n)]
        assert changed == sorted(changed) or True
        assert set(changed) == {
            "[Content_Types].xml", "xl/_rels/workbook.xml.rels", "xl/workbook.xml",
        }, changed
        assert not any(n.startswith("xl/worksheets/") for n in changed), changed


class TestProperty2NoDynamicCarrierIntroduced:
    """Property 2: 静态注入不引入动态载体（Validates: Requirements 2.5, 2.6, 2.9, 2.14）。"""

    @PBT
    @given(payload=source_and_spec())
    def test_table_part_count_and_columns_unchanged(self, payload, gate) -> None:
        source, spec = payload
        got = inject(source, spec, gate)
        assert table_part_count(got.instrumented_bytes) == table_part_count(source)
        assert columns_by_sheet(got.instrumented_bytes) == columns_by_sheet(source)

    def test_signature_cannot_express_table_or_uuid_column(self) -> None:
        """DEC-3 由类型边界保证：签名只有四项入参，写不出 Table / 隐藏列。"""
        import inspect

        params = list(inspect.signature(_xi.instrument_workbook_bytes_static_only)
                      .parameters)
        assert params == ["source", "spec", "gate", "identity_carriers"], params

    def test_mutation_proof_4_table_part_judgement_really_bites(self, gate) -> None:
        """变异证明④：人为往产物注一个 Table 部件后，「无新增 Table 部件」判据必须变红。

        结构性零（`xl/tables/` 新增 0 个）必须配变异证明 —— 否则判据写错也看不出来。
        """
        source = build_source_workbook(["Alpha"])
        spec = ExcelStaticOnlyInstrumentationSpec(
            entry_id="e", template_id="T0", template_relative_path="A/x.xlsx",
            static_regions=(region(excel_name="Alpha", managed_ref="$A$1:$C$8"),),
        )
        got = inject(source, spec, gate)
        assert table_part_count(got.instrumented_bytes) == table_part_count(source)

        # 变异：手工塞一个 Table 部件进产物 zip
        with zipfile.ZipFile(io.BytesIO(got.instrumented_bytes)) as zf:
            entries = {n: zf.read(n) for n in zf.namelist()}
        entries["xl/tables/tableMutant.xml"] = b"<table/>"
        buf = io.BytesIO()
        with zipfile.ZipFile(buf, "w", zipfile.ZIP_DEFLATED) as out:
            for name, blob in entries.items():
                out.writestr(name, blob)
        mutated = buf.getvalue()

        assert table_part_count(mutated) == table_part_count(source) + 1
        assert table_part_count(mutated) != table_part_count(source), (
            "变异证明失败：注入一个 Table 部件后判据仍为真 ⇒ 该判据是恒真装饰"
        )


class TestProperty3GtSyncKeysAndUuidSentinel:
    """Property 3: `_GT_SYNC` 键完备且 UUID 列取空串哨兵
    （Validates: Requirements 2.3, 2.4, 2.8）。"""

    @PBT
    @given(payload=source_and_spec())
    def test_keys_superset_sentinel_and_zero_runtime_binding(self, payload, gate) -> None:
        source, spec = payload
        got = inject(source, spec, gate)
        # 现算常量，禁写死键数
        assert set(_xi.REQUIRED_GT_SYNC_KEYS) <= set(got.gt_sync_pairs)
        assert got.gt_sync_pairs["GT_ROW_UUID_COLUMN"] == ""
        assert not (set(_xi.RUNTIME_BINDING_KEYS) & set(got.gt_sync_pairs))
        _xi.assert_no_runtime_binding(got.gt_sync_pairs)

    @PBT
    @given(payload=source_and_spec())
    def test_defined_names_appear_exactly_once_with_expected_ref(self, payload, gate) -> None:
        """Requirement 2.7：每个 defined_name 在 `<definedNames>` 里恰 1 次，ref 形态固定。"""
        source, spec = payload
        got = inject(source, spec, gate)
        with zipfile.ZipFile(io.BytesIO(got.instrumented_bytes)) as zf:
            wb_xml = zf.read("xl/workbook.xml").decode("utf-8")
        for r in spec.static_regions:
            assert wb_xml.count(f'<definedName name="{r.defined_name}">') == 1, r
            expected = _xi._quote_sheet_name(r.excel_name) + "!" + r.managed_ref
            assert got.defined_name_refs[r.defined_name] == expected
            assert _xi._xml_escape(expected) in wb_xml

    @PBT
    @given(payload=source_and_spec())
    def test_managed_sheet_id_is_the_real_sheet_id_of_the_defined_name_target(
        self, payload, gate
    ) -> None:
        """Requirement 2.13：`GT_MANAGED_SHEET_ID` 写 definedName 所指 sheet 的真实 sheetId。"""
        source, spec = payload
        got = inject(source, spec, gate)
        primary = spec.static_regions[0]
        with zipfile.ZipFile(io.BytesIO(source)) as zf:
            src_wb = zf.read("xl/workbook.xml").decode("utf-8")
        assert got.gt_sync_pairs["GT_MANAGED_SHEET_ID"] == _xi._sheet_id_for(
            src_wb, primary.excel_name
        )
        assert got.managed_sheet_name_at_instrumentation == primary.excel_name


class TestStaticInjectorErrorPaths:
    """两条错误路径（Requirements 2.10 / 2.11）。"""

    def test_missing_sheet_message_lists_both_missing_name_and_existing_sheets(
        self, gate
    ) -> None:
        source = build_source_workbook(["Alpha", "Beta 2"])
        spec = ExcelStaticOnlyInstrumentationSpec(
            entry_id="e", template_id="T0", template_relative_path="A/x.xlsx",
            static_regions=(region(excel_name="不存在的表", managed_ref="$A$1:$B$2"),),
        )
        with pytest.raises(InstrumentationError) as exc:
            inject(source, spec, gate)
        msg = str(exc.value)
        assert "不存在的表" in msg, msg
        # 源册现有 sheet 名清单必须一并给出 —— 只说「缺谁」定位不到「模板被换过」
        assert "Alpha" in msg and "Beta 2" in msg, msg

    def test_duplicate_defined_name_reuses_existing_static_message_shape(self, gate) -> None:
        """重名消息形态沿用既有静态支的 `Duplicate static definedName`。"""
        source = build_source_workbook(["Alpha", "Beta 2"])
        spec = ExcelStaticOnlyInstrumentationSpec(
            entry_id="e", template_id="T0", template_relative_path="A/x.xlsx",
            static_regions=(region(excel_name="Alpha", managed_ref="$A$1:$B$2"),),
        )
        once = inject(source, spec, gate)
        # 把已注入产物当源册再注一次同名 definedName：先被「已含 instrumentation 部件」挡住，
        # 故这里只留 definedName 而剥掉 `_GT_SYNC`，专门走重名路径。
        with zipfile.ZipFile(io.BytesIO(once.instrumented_bytes)) as zf:
            entries = {n: zf.read(n) for n in zf.namelist()
                       if n != "xl/worksheets/sheetGtSync.xml"}
        wb = entries["xl/workbook.xml"].decode("utf-8")
        wb = re.sub(r'<sheet[^>]*name="_GT_SYNC"[^>]*/>', "", wb)
        entries["xl/workbook.xml"] = wb.encode("utf-8")
        rels = entries["xl/_rels/workbook.xml.rels"].decode("utf-8")
        entries["xl/_rels/workbook.xml.rels"] = re.sub(
            r'<Relationship Id="rIdGTSYNC".*?/>', "", rels
        ).encode("utf-8")
        ct = entries["[Content_Types].xml"].decode("utf-8")
        entries["[Content_Types].xml"] = re.sub(
            r'<Override PartName="/xl/worksheets/sheetGtSync\.xml".*?/>', "", ct
        ).encode("utf-8")
        buf = io.BytesIO()
        with zipfile.ZipFile(buf, "w", zipfile.ZIP_DEFLATED) as out:
            for name, blob in entries.items():
                out.writestr(name, blob)

        with pytest.raises(InstrumentationError) as exc:
            inject(buf.getvalue(), spec, gate)
        assert "Duplicate static definedName" in str(exc.value), str(exc.value)

    def test_empty_identity_carriers_is_rejected(self, gate) -> None:
        source = build_source_workbook(["Alpha"])
        spec = ExcelStaticOnlyInstrumentationSpec(
            entry_id="e", template_id="T0", template_relative_path="A/x.xlsx",
            static_regions=(region(excel_name="Alpha", managed_ref="$A$1:$B$2"),),
        )
        with pytest.raises(InstrumentationError) as exc:
            _xi.instrument_workbook_bytes_static_only(
                source, spec, gate=gate, identity_carriers=()
            )
        assert "identity_carriers" in str(exc.value)

    def test_reinjection_of_instrumented_artifact_is_rejected(self, gate) -> None:
        source = build_source_workbook(["Alpha"])
        spec = ExcelStaticOnlyInstrumentationSpec(
            entry_id="e", template_id="T0", template_relative_path="A/x.xlsx",
            static_regions=(region(excel_name="Alpha", managed_ref="$A$1:$B$2"),),
        )
        once = inject(source, spec, gate)
        with pytest.raises(InstrumentationError) as exc:
            inject(once.instrumented_bytes, spec, gate)
        assert "已含 instrumentation 部件" in str(exc.value)


class TestMultiInjectorWasNotTouched:
    """对照组：`instrument_workbook_bytes_multi` 一字不改（Requirement 2.2）。"""

    def test_multi_injector_still_rejects_empty_specs(self, gate) -> None:
        with pytest.raises(InstrumentationError) as exc:
            _xi.instrument_workbook_bytes_multi(b"", (), gate=gate)
        assert "specs 不得为空" in str(exc.value)

    def test_multi_injector_still_hardcodes_all_six_primary_row_attrs(self) -> None:
        import inspect

        src = inspect.getsource(_xi.instrument_workbook_bytes_multi)
        for attr in ("primary.footer_row", "primary.uuid_col", "primary.table_name",
                     "primary.table_ref", "primary.first_data_row",
                     "primary.last_data_row"):
            assert attr in src, f"`instrument_workbook_bytes_multi` 不再硬取 {attr}"


# ═══════════════════════════════════════════════════════════════════════════
# §3 Static_Payload_Builder 与 digest 不变式（Requirements 3.1 ~ 3.11）
# ═══════════════════════════════════════════════════════════════════════════


@pytest.fixture(scope="module")
def a51() -> Any:
    from app.services.workpaper_sync import phase5_a51_cashflow_audit as mod

    return mod


class TestStaticPayloadBuilderShape:
    """载体清单来源、空集合语义、键的有无（Requirements 3.1 ~ 3.7）。"""

    def test_signature_is_exactly_the_six_declared_params(self) -> None:
        import inspect

        params = list(inspect.signature(
            _xi.build_static_only_instrumentation_payload).parameters)
        assert params == [
            "spec", "template_definition_sha256", "template_sha256",
            "gate", "identity_carriers", "identity_anchors",
        ], params

    def test_identity_carriers_come_from_the_argument_not_from_the_gate(
        self, a51: Any, gate: Any
    ) -> None:
        """Requirement 3.3：A5-1 的载体清单项数 < gate 放行项数 ⇒ 按 gate 推导即过度声明。"""
        payload = a51.build_instrumentation_payload()
        assert payload["identity_carriers"] == list(a51.IDENTITY_CARRIERS)
        assert len(a51.IDENTITY_CARRIERS) < len(gate.allowed_carriers), (
            "分母前提失效：载体清单已不小于 gate 放行集，本判据须重算"
        )
        assert payload["identity_carriers"] != sorted(gate.allowed_carriers)
        assert payload["identity_anchors"] == list(a51.IDENTITY_ANCHORS)

    def test_managed_sheets_empty_and_static_sheets_count_matches_spec(
        self, a51: Any
    ) -> None:
        payload = a51.build_instrumentation_payload()
        spec = a51.static_only_instrumentation_spec()
        assert payload["managed_sheets"] == []
        assert len(payload["static_sheets"]) == len(spec.static_regions)

    def test_payload_has_neither_row_uuid_disposition_nor_transposed_sheets(
        self, a51: Any
    ) -> None:
        payload = a51.build_instrumentation_payload()
        assert "row_uuid_disposition" not in payload
        assert "transposed_sheets" not in payload

    def test_hidden_metadata_keys_reuse_the_platform_constant_verbatim(
        self, a51: Any
    ) -> None:
        """Requirement 3.6：原样复用 `REQUIRED_GT_SYNC_KEYS`，禁新建静态子集。"""
        payload = a51.build_instrumentation_payload()
        assert payload["hidden_metadata_sheet"]["keys"] == list(
            _xi.REQUIRED_GT_SYNC_KEYS
        )

    def test_cell_geometry_is_the_static_form_block(self, a51: Any) -> None:
        geo = a51.build_instrumentation_payload()["cell_geometry"]
        assert set(geo) == {"anchor_role", "note", "sheets"}, sorted(geo)
        assert "table_ref" not in geo and "footer_row" not in geo
        assert "managed_range" not in geo          # 顶层无（只在 sheets[] 里）
        for sheet in geo["sheets"]:
            assert set(sheet) == {"sheet_key", "managed_range", "region_kind"}
            assert sheet["region_kind"] == "static"

    def test_payload_passes_both_platform_validators(self, a51: Any, gate: Any) -> None:
        from app.services.workpaper_sync.definitions import (
            validate_instrumentation_payload,
        )

        payload = a51.build_instrumentation_payload()
        validate_instrumentation_payload(payload)
        _xi._assert_no_forbidden_anchor_declared(payload, gate=gate)

    @pytest.mark.parametrize("bad", ["", "zz", "abc", "0" * 63])
    def test_builder_rejects_non_digest_inputs(self, a51: Any, gate: Any, bad: str) -> None:
        with pytest.raises(InstrumentationError):
            _xi.build_static_only_instrumentation_payload(
                spec=a51.static_only_instrumentation_spec(),
                template_definition_sha256=bad,
                template_sha256=a51.TEMPLATE_SHA256,
                gate=gate,
                identity_carriers=a51.IDENTITY_CARRIERS,
                identity_anchors=a51.IDENTITY_ANCHORS,
            )

    def test_builder_rejects_empty_carrier_or_anchor_lists(self, a51: Any, gate: Any) -> None:
        from app.services.workpaper_sync.definitions import canonical_digest

        common = dict(
            spec=a51.static_only_instrumentation_spec(),
            template_definition_sha256=canonical_digest(a51.template_definition_payload()),
            template_sha256=a51.TEMPLATE_SHA256,
            gate=gate,
        )
        with pytest.raises(InstrumentationError):
            _xi.build_static_only_instrumentation_payload(
                identity_carriers=(), identity_anchors=a51.IDENTITY_ANCHORS, **common
            )
        with pytest.raises(InstrumentationError):
            _xi.build_static_only_instrumentation_payload(
                identity_carriers=a51.IDENTITY_CARRIERS, identity_anchors=(), **common
            )


class TestGenericPayloadBuilderWasNotRelaxed:
    """对照组：`build_instrumentation_payload_for_sheets` 保持改动前形态（Requirement 3.2）。"""

    def test_generic_builder_still_rejects_empty_specs(self, gate: Any) -> None:
        with pytest.raises(InstrumentationError) as exc:
            _xi.build_instrumentation_payload_for_sheets(
                specs=(), template_definition_sha256="a" * 64,
                template_sha256="b" * 64, gate=gate,
            )
        assert "specs 不得为空" in str(exc.value)

    def test_generic_builder_still_derives_carriers_from_the_gate(self) -> None:
        import inspect

        src = inspect.getsource(_xi.build_instrumentation_payload_for_sheets)
        assert "list(sorted(gate.allowed_carriers))" in src, (
            "通用构建器的 identity_carriers 取值逻辑被改动 —— 它必须逐条保留"
        )


class TestDigestInvariantIsTheLaneAAcceptanceGate:
    """A-4 digest 不变式 —— **EXAMPLE 不是属性**。

    入参是**单一冻结值**（canary 的那一组），写成 `hypothesis` 随机化测试即误用：
    随机 spec 的 digest 本来就该不同，随机化反而把判据冲淡。

    🔴 **三条禁止**：不相等时只能修 Static_Payload_Builder；改契约 / 改守卫常量 /
    改真库 artifact 去迁就 = 不合格，须回退。
    """

    def test_platform_builder_digest_equals_the_frozen_canary_digest(
        self, a51: Any, gate: Any
    ) -> None:
        from app.services.workpaper_sync.definitions import canonical_digest

        # 平台构建器：以 canary 的固定入参直调（不经 provider 门面）
        platform = _xi.build_static_only_instrumentation_payload(
            spec=a51.static_only_instrumentation_spec(),
            template_definition_sha256=canonical_digest(a51.template_definition_payload()),
            template_sha256=a51.TEMPLATE_SHA256,
            gate=gate,
            identity_carriers=a51.IDENTITY_CARRIERS,
            identity_anchors=a51.IDENTITY_ANCHORS,
        )
        # provider 门面：委派同一入口
        via_provider = a51.build_instrumentation_payload()

        if canonical_digest(platform) != canonical_digest(via_provider):
            # 不相等时输出 key 级差异（只说「不等」定位不到哪一项漂了）
            keys = sorted(set(platform) | set(via_provider))
            diff = [
                (k, platform.get(k, "<缺>"), via_provider.get(k, "<缺>"))
                for k in keys
                if platform.get(k, "<缺>") != via_provider.get(k, "<缺>")
            ]
            pytest.fail(f"digest 不相等，key 级差异: {diff}")

        # ── 三方同值（判据以常量为锚，禁在判据文本里写死全值）──────────
        from test_a_entry_connection_blockers import (  # noqa: PLC0415
            CANARY_INSTRUMENTATION_DEFINITION_SHA256 as GUARD_CONST,
        )

        contract_path = (
            _BACKEND / "data" / "workpaper_sync_contracts" / "a51.cashflow_audit.json"
        )
        import json

        contract = json.loads(contract_path.read_bytes().decode("utf-8"))
        assert canonical_digest(platform) == GUARD_CONST
        assert contract["instrumentation_definition_sha256"] == GUARD_CONST
        assert canonical_digest(a51.instrumentation_definition_payload()) == GUARD_CONST

    def test_frozen_note_constant_is_the_single_source_of_the_geometry_note(
        self, a51: Any
    ) -> None:
        """note 文字是冻结值：改一个字就改 digest ⇒ 它必须只有一处真源。"""
        payload = a51.build_instrumentation_payload()
        assert payload["cell_geometry"]["note"] == _xi._STATIC_CELL_GEOMETRY_NOTE
        import inspect

        provider_src = inspect.getsource(a51)
        assert "仅描述注入时刻的几何以便复现 instrumentation" not in provider_src, (
            "provider 又抄了一份 note 文字 —— 两份必漂"
        )

    def test_provider_no_longer_assembles_the_payload_itself(self, a51: Any) -> None:
        """Requirement 3.10：provider 收缩为门面（形态由平台表达）。"""
        import inspect

        src = inspect.getsource(a51.build_instrumentation_payload)
        assert "build_static_only_instrumentation_payload" in src
        for token in ('"schema_version"', '"managed_sheets"', '"cell_geometry"',
                      '"hidden_metadata_sheet"'):
            assert token not in src, f"门面里仍在自组装 {token}"

    def test_stale_contract_denominator_comment_was_corrected(self, a51: Any) -> None:
        """Requirement 3.11：过期分母注释已改成现算口径，且不写死任何分母数字。"""
        raw = Path(a51.__file__).read_bytes().decode("utf-8")
        assert "60 份契约" not in raw and "实测 60" not in raw
        doc = a51.build_instrumentation_payload.__doc__ or ""
        assert "现算" in doc, doc
        assert not re.search(r"\d{2,}\s*份", doc), (
            f"门面 docstring 里仍写死了分母数字: {doc!r}"
        )


# ═══════════════════════════════════════════════════════════════════════════
# §4 Substrate 组装的第三条分派臂（Requirements 4.1 ~ 4.5）
# ═══════════════════════════════════════════════════════════════════════════

ARM_MULTI, ARM_SINGLE, ARM_STATIC, ARM_CRASH = "multi", "single", "static_only", "crash"


def arm_before(provider: Any) -> str:
    """改动**前**的臂选择规则（照搬旧源码：`if callable(instrumentation_specs) else ...`）。

    旧第二臂是无条件 `else`，对没有 `instrumentation_spec` 的 provider 抛裸
    `AttributeError` ⇒ 这里用 `ARM_CRASH` 表达那个形态。
    """
    if callable(getattr(provider, "instrumentation_specs", None)):
        return ARM_MULTI
    if callable(getattr(provider, "instrumentation_spec", None)):
        return ARM_SINGLE
    return ARM_CRASH


def arm_after(provider: Any) -> str:
    """改动**后**的臂选择规则（与生产代码的 if / elif / else 同构）。"""
    if callable(getattr(provider, "instrumentation_specs", None)):
        return ARM_MULTI
    if callable(getattr(provider, "instrumentation_spec", None)):
        return ARM_SINGLE
    if callable(getattr(provider, "static_only_instrumentation_spec", None)):
        return ARM_STATIC
    return ARM_CRASH


@pytest.fixture(scope="module")
def delivered_providers() -> list[tuple[str, Any]]:
    """交付台账里登记的全部 provider（现算分母，禁写死成员）。"""
    import importlib

    from app.services.workpaper_sync.adapters.registry import (
        DELIVERED_PER_ENTRY_CONTRACTS,
    )

    out: list[tuple[str, Any]] = []
    for row in DELIVERED_PER_ENTRY_CONTRACTS:
        module_path = str(row.get("provider_module") or "")
        if not module_path:
            continue
        try:
            out.append((str(row.get("entry_id") or ""), importlib.import_module(module_path)))
        except Exception:  # noqa: BLE001 — 不可导入的 provider 不在本判据分母内
            continue
    assert out, "交付台账 provider 分母为空 —— 判据失去分母"
    return out


class TestThirdDispatchArm:
    """第三臂追加在既有两臂之后，且既有 provider 零走向变更。"""

    def test_dispatch_order_in_source_is_specs_then_spec_then_static_only(self) -> None:
        import inspect

        from app.services.workpaper_sync import projection_first_publication as pfp

        src = inspect.getsource(pfp.stage_instrumented_substrate)
        i_specs = src.index('"instrumentation_specs"')
        i_spec = src.index('"instrumentation_spec"')
        i_static = src.index('"static_only_instrumentation_spec"')
        # 分派判据的先后：`callable(specs_fn)` → `callable(spec_fn)` → 第三臂
        j_specs = src.index("if callable(specs_fn):")
        j_spec = src.index("elif callable(spec_fn):")
        j_static = src.index("instrument_workbook_bytes_static_only(")
        assert j_specs < j_spec < j_static, (j_specs, j_spec, j_static)
        assert min(i_specs, i_spec, i_static) >= 0

    def test_no_existing_provider_changes_arm(
        self, delivered_providers: list[tuple[str, Any]]
    ) -> None:
        """现算分母内，改动前后命中的臂逐个相同（canary 除外，它原本是 ARM_CRASH）。"""
        moved: list[tuple[str, str, str]] = []
        canaries: list[str] = []
        for entry_id, mod in delivered_providers:
            before, after = arm_before(mod), arm_after(mod)
            if before == after:
                continue
            if before == ARM_CRASH and after == ARM_STATIC:
                canaries.append(entry_id)       # 这正是本 lane 要接通的形态
                continue
            moved.append((entry_id, before, after))
        assert not moved, f"既有 provider 走向被改变: {moved}"
        assert canaries, (
            "没有任何 provider 从「裸 AttributeError」变成第三臂 ⇒ 本 lane 没接通任何 entry"
        )

    def test_canary_hits_the_third_arm(self, a51: Any) -> None:
        assert arm_before(a51) == ARM_CRASH
        assert arm_after(a51) == ARM_STATIC

    def test_ooxml_security_gate_still_precedes_staged_artifact(self) -> None:
        """Requirement 4.5：安全门仍排在产出 staged artifact **之前**。"""
        import inspect

        from app.services.workpaper_sync import projection_first_publication as pfp

        src = inspect.getsource(pfp.stage_instrumented_substrate)
        assert src.index("validate_ooxml_artifact(") < src.index("probe_path.replace(")


class TestDec3StructuralGuard:
    """fail-closed 单向蕴含守卫（Requirements 4.3 / 4.4）。"""

    def _stage(self, provider: Any, tmp_path: Path):
        from app.services.workpaper_sync import projection_first_publication as pfp

        original = pfp._provider_for
        pfp._provider_for = lambda entry_id: provider          # type: ignore[assignment]
        try:
            return pfp.stage_instrumented_substrate(
                entry_id="probe/entry", staging_dir=tmp_path, contract=object()
            )
        finally:
            pfp._provider_for = original                        # type: ignore[assignment]

    def test_both_static_only_and_row_spec_is_fail_closed(
        self, a51: Any, tmp_path: Path
    ) -> None:
        from app.services.workpaper_sync.projection_first_publication import (
            SubstrateStagingError,
        )

        class Hybrid:
            read_authoritative_template = staticmethod(a51.read_authoritative_template)
            excel_carrier_gate = staticmethod(a51.excel_carrier_gate)
            static_only_instrumentation_spec = staticmethod(
                a51.static_only_instrumentation_spec
            )
            IDENTITY_CARRIERS = a51.IDENTITY_CARRIERS

            @staticmethod
            def instrumentation_spec():        # 行表 spec —— DEC-3 明禁的形态
                raise AssertionError("守卫必须先于分派触发，不得走到这里")

        with pytest.raises(SubstrateStagingError) as exc:
            self._stage(Hybrid, tmp_path)
        msg = str(exc.value)
        assert "static_only_instrumentation_spec" in msg and "DEC-3" in msg, msg

    def test_guard_fires_even_when_arm_one_would_have_won(
        self, a51: Any, tmp_path: Path
    ) -> None:
        """守卫**先于**分派：同时有 `instrumentation_specs` 时也必须被挡住。

        若守卫写在「进第三臂时才查」，这个 provider 会先命中第一臂 ⇒ 永远查不到。
        """
        from app.services.workpaper_sync.projection_first_publication import (
            SubstrateStagingError,
        )

        class HybridMulti:
            read_authoritative_template = staticmethod(a51.read_authoritative_template)
            excel_carrier_gate = staticmethod(a51.excel_carrier_gate)
            static_only_instrumentation_spec = staticmethod(
                a51.static_only_instrumentation_spec
            )
            IDENTITY_CARRIERS = a51.IDENTITY_CARRIERS

            @staticmethod
            def instrumentation_specs():
                raise AssertionError("守卫必须先于第一臂触发")

        with pytest.raises(SubstrateStagingError) as exc:
            self._stage(HybridMulti, tmp_path)
        assert "instrumentation_specs" in str(exc.value)

    def test_negative_dual_row_entry_provider_is_not_flagged(
        self, a51: Any, tmp_path: Path
    ) -> None:
        """负向：同时暴露 `instrumentation_spec` 与 `instrumentation_specs` 的**既有形态**
        provider 不得被打红 —— 那是双入口回落，不是 DEC-3 违规。

        这一条证明守卫用的是单向蕴含，而不是「三者恰暴露其一」的统一断言。
        """
        sentinel = RuntimeError("已进入第一臂（说明未被误判）")

        class DualRowEntry:
            read_authoritative_template = staticmethod(a51.read_authoritative_template)
            excel_carrier_gate = staticmethod(a51.excel_carrier_gate)

            @staticmethod
            def instrumentation_specs():
                raise sentinel

            @staticmethod
            def instrumentation_spec():
                raise AssertionError("不应走到第二臂")

        from app.services.workpaper_sync.projection_first_publication import (
            SubstrateStagingError,
        )

        with pytest.raises(SubstrateStagingError) as exc:
            self._stage(DualRowEntry, tmp_path)
        msg = str(exc.value)
        assert "已进入第一臂" in msg, msg
        assert "DEC-3" not in msg, f"双入口回落形态被误判成 DEC-3 违规: {msg}"

    def test_existing_dual_entry_providers_really_exist_in_the_denominator(
        self, delivered_providers: list[tuple[str, Any]]
    ) -> None:
        """变异证明：上一条的「既有形态」不是假想 —— 现算分母里真有双入口 provider。

        若这里现算为 0，说明「统一断言会打红既有 entry」这个理由已失效，
        上一条负向判据须重算（而不是继续照抄）。
        """
        dual = [
            entry_id
            for entry_id, mod in delivered_providers
            if callable(getattr(mod, "instrumentation_specs", None))
            and callable(getattr(mod, "instrumentation_spec", None))
        ]
        assert dual, (
            "现算分母内无双入口 provider ⇒ 单向蕴含的理由须重算"
        )

    def test_static_provider_without_identity_carriers_is_rejected(
        self, a51: Any, tmp_path: Path
    ) -> None:
        from app.services.workpaper_sync.projection_first_publication import (
            SubstrateStagingError,
        )

        class NoCarriers:
            read_authoritative_template = staticmethod(a51.read_authoritative_template)
            excel_carrier_gate = staticmethod(a51.excel_carrier_gate)
            static_only_instrumentation_spec = staticmethod(
                a51.static_only_instrumentation_spec
            )

        with pytest.raises(SubstrateStagingError) as exc:
            self._stage(NoCarriers, tmp_path)
        assert "IDENTITY_CARRIERS" in str(exc.value)


# ═══════════════════════════════════════════════════════════════════════════
# §5 第 6 处阻塞：观测清册（Requirements 6.1 ~ 6.8 / Property 6）
# ═══════════════════════════════════════════════════════════════════════════

from app.services.workpaper_sync import published_identity_observer as OBS  # noqa: E402
from app.services.workpaper_sync.published_identity_observer import (  # noqa: E402
    StaticIdentityInventory,
)


@pytest.fixture(scope="module")
def canary_instrumented(a51: Any, gate: Any) -> bytes:
    return _xi.instrument_workbook_bytes_static_only(
        a51.read_authoritative_template(),
        a51.static_only_instrumentation_spec(),
        gate=gate,
        identity_carriers=a51.IDENTITY_CARRIERS,
    ).instrumented_bytes


@pytest.fixture(scope="module")
def canary_contract(a51: Any) -> Any:
    return a51.load_contract_from_disk()


@pytest.fixture(scope="module")
def canary_anchors(a51: Any) -> list[dict[str, str]]:
    anchors = OBS._frozen_sheet_anchors(a51.build_instrumentation_payload())
    assert anchors, "静态锚点为空 —— 判据失去分母"
    assert all(a.get("region_kind") == "static" for a in anchors), anchors
    return anchors


def collect(data: bytes, contract: Any, anchors: list[dict[str, str]]):
    return OBS.collect_workbook_structure(
        data=data, contract=contract, sheet_anchors=anchors
    )


class TestStaticIdentityInventoryShape:
    """清册的成员面（Requirement 6.2 + 消费方审计结论 6.6）。"""

    def test_two_mandated_members(self) -> None:
        inv = StaticIdentityInventory(regions=(("k", "GT_A", "Sheet1"),))
        assert inv.row_uuids == ()
        assert inv.inventory_digest_input == {
            "kind": "static_region", "regions": [["k", "GT_A", "Sheet1"]]
        }

    def test_members_required_by_audited_consumers_are_present(self) -> None:
        """消费方审计（Task 7.2）现算结论：`excel_extract._assert_static_anchor_retained`
        用 `defined_names` 与 `hidden_sheet_present` ⇒ 静态清册必须真实提供它们，
        否则 `AttributeError` 只是从 `_observe_workbook` 搬到 `excel_extract`。"""
        inv = StaticIdentityInventory(
            regions=(("k", "GT_A", "Sheet1"), ("k2", "GT_B", "Sheet2")),
            hidden_sheet_present=True, hidden_sheet_is_hidden=True,
            excluded_from_business_enumeration=True,
        )
        assert inv.defined_names == ("GT_A", "GT_B")
        assert inv.hidden_sheet_present is True
        assert inv.is_static_region_inventory is True
        assert inv.duplicate_row_uuids == () and inv.empty_row_uuids == ()

    def test_row_table_only_members_are_deliberately_absent(self) -> None:
        """🔴 行表事实**不得伪造** —— 伪造即 DEC-3 明禁的「注退化动态表当载体」。"""
        inv = StaticIdentityInventory(regions=(("k", "GT_A", "Sheet1"),))
        for member in ("table_present", "table_ref", "resolved_sheet_by",
                       "uuid_column_hidden"):
            assert not hasattr(inv, member), f"静态清册伪造了行表事实 {member!r}"

    def test_digest_input_excludes_hidden_sheet_facts(self) -> None:
        """Requirement 6.2 钉死 digest 形态：只有 kind + regions。"""
        a = StaticIdentityInventory(regions=(("k", "GT_A", "S"),))
        b = StaticIdentityInventory(regions=(("k", "GT_A", "S"),),
                                    hidden_sheet_present=True)
        assert a.inventory_digest_input == b.inventory_digest_input


class TestProperty6NoBareExceptionUnderAllStaticAnchors:
    """Property 6: 全静态锚点下观测无裸异常
    （Validates: Requirements 6.1, 6.5）。"""

    @PBT
    @given(region_count=st.integers(min_value=1, max_value=2))
    def test_single_and_multi_region_both_yield_a_real_inventory(
        self, region_count: int, a51: Any, canary_instrumented: bytes,
        canary_contract: Any, canary_anchors: list[dict[str, str]]
    ) -> None:
        """单区与多区两形态都要覆盖。

        🔴 单区形态**必须配 focused 契约**：锚点子集配完整契约时，未被观测的那张 sheet
        拿不到 physical 映射 ⇒ `assert_no_structure_drift` 抛 `ContractDriftError`，
        那是**正确**行为（不是本属性要抓的缺陷）。照抄「锚点切片 + 完整契约」会把一个
        正确判据当成缺陷。
        """
        from app.services.workpaper_sync.contracts import parse_contract

        subset = canary_anchors[:region_count]
        keys = {a["sheet_key"] for a in subset}
        payload = dict(canary_contract.canonical_payload)
        payload["sheets"] = [s for s in payload["sheets"] if s["sheet_key"] in keys]
        contract = parse_contract(payload, adapter_id=canary_contract.contract_id)

        try:
            _fp, _physical, inventory, _structure = collect(
                canary_instrumented, contract, subset
            )
        except AttributeError as exc:            # 这正是本任务要消除的形态
            pytest.fail(f"全静态锚点下抛了裸 AttributeError: {exc}")
        except (OBS.FrozenChildUnusableError, OBS.ObservedIdentityDriftError):
            return                               # 允许的失败形式（显式异常类型）
        assert inventory is not None, "第 3 返回值仍为 None ⇒ 第 6 处阻塞未解除"
        assert isinstance(inventory, StaticIdentityInventory)
        assert inventory.row_uuids == ()
        assert len(inventory.regions) == region_count

    def test_all_static_anchors_produce_static_inventory(
        self, canary_instrumented: bytes, canary_contract: Any,
        canary_anchors: list[dict[str, str]]
    ) -> None:
        _fp, physical, inventory, _st = collect(
            canary_instrumented, canary_contract, canary_anchors
        )
        assert isinstance(inventory, StaticIdentityInventory)
        assert len(inventory.regions) == len(canary_anchors)
        # 三元组的第 3 项是**物理** sheet 名（由 definedName 反读得出）
        for sheet_key, defined_name, physical_sheet in inventory.regions:
            assert physical[sheet_key] == physical_sheet
            assert defined_name in inventory.defined_names
        # 隐藏 metadata sheet 三项事实是真实观测值
        assert inventory.hidden_sheet_present is True
        assert inventory.hidden_sheet_is_hidden is True
        assert inventory.excluded_from_business_enumeration is True

    def test_downstream_gate_accepts_the_static_inventory(
        self, a51: Any, canary_instrumented: bytes, canary_contract: Any,
        canary_anchors: list[dict[str, str]]
    ) -> None:
        """消费方审计的**闭环证明**：下游 gate 真能吃下静态清册（不是只看类型对）。"""
        from app.services.workpaper_sync.excel_entry_gate import (
            assert_identity_inventory_usable,
        )

        _fp, _physical, inventory, _st = collect(
            canary_instrumented, canary_contract, canary_anchors
        )
        assert_identity_inventory_usable(
            inventory, contract=canary_contract, entry_id=a51.ENTRY_ID
        )

    def test_static_arm_returns_before_touching_row_table_members(self) -> None:
        """静态分派臂必须在触及行表成员**之前** return（现算偏移，禁写死）。"""
        import inspect

        from app.services.workpaper_sync.excel_entry_gate import (
            assert_identity_inventory_usable,
        )

        src = inspect.getsource(assert_identity_inventory_usable)
        i_static = src.index("is_static_region_inventory")
        i_return = src.index("        return", i_static)
        first_row_member = min(
            src.index(f"inventory.{m}")
            for m in ("table_present", "table_ref", "resolved_sheet_by",
                      "uuid_column_hidden")
            if f"inventory.{m}" in src
        )
        assert i_static < i_return < first_row_member, (
            i_static, i_return, first_row_member
        )

    def test_observe_workbook_digest_line_is_untouched(self) -> None:
        """Requirement 6.3：处置方式是**消除 None**，不是给那行加 None 特判。"""
        import inspect

        src = inspect.getsource(OBS.PublishedIdentityObserver._observe_workbook)
        assert "canonical_digest(inventory.inventory_digest_input)" in src
        assert "inventory is None" not in src
        assert "Optional" not in src

    def test_observe_workbook_runs_before_build_identity_binding(self) -> None:
        """Requirement 6.8：第 6 处在执行顺序上**早于**第 5 处（位置值现算，禁写死）。

        这是「Task 7 必须整体先于 Task 8」的承重实证：`_observe_workbook` 先跑，它未修时
        纯静态 entry 会先撞裸 `AttributeError`，**到不了** `_build_identity_binding` 的
        `row_identity` 判据 ⇒「静态主 binding 已支持」**不能**作为第 6 处已完成的证据。

        两个口径（字符偏移 + 相对行序）互证，单口径可能因某次改写而口误。
        """
        import inspect

        src = inspect.getsource(OBS.PublishedIdentityObserver.observe)
        i_wb, i_bd = src.find("_observe_workbook"), src.find("_build_identity_binding")
        assert i_wb >= 0 and i_bd >= 0, (i_wb, i_bd)
        assert i_wb < i_bd, f"字符偏移口径：{i_wb} 应 < {i_bd}"
        lines = src.split("\n")
        l_wb = next(i for i, ln in enumerate(lines) if "_observe_workbook" in ln)
        l_bd = next(i for i, ln in enumerate(lines) if "_build_identity_binding" in ln)
        assert l_wb < l_bd, f"行序口径：{l_wb} 应 < {l_bd}"

    def test_dynamic_path_still_returns_the_original_primary_inventory(self) -> None:
        """Requirement 6.7：含动态 Excel-Table 锚点时第 3 返回值仍是改动前的对象。"""
        import inspect

        src = inspect.getsource(OBS.collect_workbook_structure)
        i_dynamic = src.index("if primary_inventory is None:")
        i_static_fallback = src.index("if primary_inventory is None and static_observed:")
        assert i_dynamic < i_static_fallback, (
            "静态兜底必须排在动态赋值**之后** —— 否则会抢掉动态清册"
        )
        assert src.count("primary_inventory = inventory") == 1


def rewrite_workbook_xml(data: bytes, fn) -> bytes:
    """对 `xl/workbook.xml` 做定点改写并重打包（其余部件字节不动）。"""
    with zipfile.ZipFile(io.BytesIO(data)) as zf:
        entries = {n: zf.read(n) for n in zf.namelist()}
    entries["xl/workbook.xml"] = fn(
        entries["xl/workbook.xml"].decode("utf-8")
    ).encode("utf-8")
    buf = io.BytesIO()
    with zipfile.ZipFile(buf, "w", zipfile.ZIP_DEFLATED) as out:
        for name, blob in entries.items():
            out.writestr(name, blob)
    return buf.getvalue()


class TestDefinedNameDriftKeepsTheDigestMeaningful:
    """Requirement 6.4：definedName 被删 / 改名 / 改指向另一张 sheet 时该 digest 有反应。

    这一组是「静态清册的 digest 是**真实观测值**而不是补出来的假值」的证明。
    少了它，`identity_inventory_sha256` 对纯静态 entry 就成了一个恒定装饰。
    """

    def _digest(self, data: bytes, contract: Any, anchors: list[dict[str, str]]) -> str:
        from app.services.workpaper_sync.definitions import canonical_digest

        _fp, _physical, inventory, _st = collect(data, contract, anchors)
        return canonical_digest(inventory.inventory_digest_input)

    def test_baseline_digest_is_derived_from_the_observed_regions(
        self, canary_instrumented: bytes, canary_contract: Any,
        canary_anchors: list[dict[str, str]]
    ) -> None:
        """先钉死「它是观测值」：digest 等于对观测三元组现算的 digest。"""
        from app.services.workpaper_sync.definitions import canonical_digest

        _fp, physical, inventory, _st = collect(
            canary_instrumented, canary_contract, canary_anchors
        )
        expected = canonical_digest({
            "kind": "static_region",
            "regions": [[a["sheet_key"], a["defined_name"], physical[a["sheet_key"]]]
                        for a in canary_anchors],
        })
        assert canonical_digest(inventory.inventory_digest_input) == expected

    def test_deleted_defined_name_fails_loudly(
        self, canary_instrumented: bytes, canary_contract: Any,
        canary_anchors: list[dict[str, str]]
    ) -> None:
        """① 删除：比「digest 变了」更强的信号 —— 显式异常，且**不是**裸 AttributeError。"""
        victim = canary_anchors[0]["defined_name"]
        broken = rewrite_workbook_xml(
            canary_instrumented,
            lambda xml: re.sub(
                rf'<definedName name="{re.escape(victim)}">[^<]*</definedName>', "", xml
            ),
        )
        with pytest.raises(OBS.ObservedIdentityDriftError) as exc:
            collect(broken, canary_contract, canary_anchors)
        assert victim in str(exc.value)

    def test_renamed_defined_name_fails_loudly(
        self, canary_instrumented: bytes, canary_contract: Any,
        canary_anchors: list[dict[str, str]]
    ) -> None:
        """② 改名：冻结锚点按旧名反读不到 ⇒ 显式漂移异常。"""
        victim = canary_anchors[0]["defined_name"]
        renamed = rewrite_workbook_xml(
            canary_instrumented,
            lambda xml: xml.replace(f'name="{victim}"', f'name="{victim}_RENAMED"'),
        )
        with pytest.raises(OBS.ObservedIdentityDriftError) as exc:
            collect(renamed, canary_contract, canary_anchors)
        assert victim in str(exc.value)

    def test_renamed_defined_name_with_updated_anchor_changes_the_digest(
        self, canary_instrumented: bytes, canary_contract: Any,
        canary_anchors: list[dict[str, str]]
    ) -> None:
        """② 之续：改名**且**同步更新锚点时，digest 必须变（名字进了 digest 输入）。"""
        base = self._digest(canary_instrumented, canary_contract, canary_anchors)
        victim = canary_anchors[0]["defined_name"]
        renamed = rewrite_workbook_xml(
            canary_instrumented,
            lambda xml: xml.replace(f'name="{victim}"', f'name="{victim}_V2"'),
        )
        new_anchors = [dict(a) for a in canary_anchors]
        new_anchors[0]["defined_name"] = f"{victim}_V2"
        assert self._digest(renamed, canary_contract, new_anchors) != base

    def test_defined_name_pointing_at_a_renamed_sheet_changes_the_digest(
        self, canary_instrumented: bytes, canary_contract: Any,
        canary_anchors: list[dict[str, str]]
    ) -> None:
        """③ 改指向：definedName 指到**另一个名字**的 sheet 时 digest 必须变。

        物理 sheet 名是 digest 输入的第 3 分量；它变了 digest 不变，就说明 digest
        没在观测物理面（那正是「补出来的假值」的形态）。
        """
        base = self._digest(canary_instrumented, canary_contract, canary_anchors)
        _fp, physical, _inv, _st = collect(
            canary_instrumented, canary_contract, canary_anchors
        )
        old_sheet = physical[canary_anchors[0]["sheet_key"]]
        new_sheet = f"{old_sheet}_MOVED"
        moved = rewrite_workbook_xml(
            canary_instrumented,
            lambda xml: xml.replace(
                f'name="{_xi._xml_escape(old_sheet)}"',
                f'name="{_xi._xml_escape(new_sheet)}"',
            ).replace(
                _xi._xml_escape(_xi._quote_sheet_name(old_sheet)),
                _xi._xml_escape(_xi._quote_sheet_name(new_sheet)),
            ),
        )
        _fp2, physical2, _i2, _s2 = collect(moved, canary_contract, canary_anchors)
        assert physical2[canary_anchors[0]["sheet_key"]] == new_sheet
        assert self._digest(moved, canary_contract, canary_anchors) != base

    def test_digest_is_sensitive_to_all_three_triple_components(self) -> None:
        """单元级兜底：三元组的**每一个**分量都进 digest（禁只对其中一两个敏感）。"""
        from app.services.workpaper_sync.definitions import canonical_digest

        base = StaticIdentityInventory(regions=(("k", "GT_A", "Sheet1"),))
        variants = (
            StaticIdentityInventory(regions=(("k2", "GT_A", "Sheet1"),)),
            StaticIdentityInventory(regions=(("k", "GT_B", "Sheet1"),)),
            StaticIdentityInventory(regions=(("k", "GT_A", "Sheet2"),)),
        )
        seen = {canonical_digest(base.inventory_digest_input)}
        for v in variants:
            d = canonical_digest(v.inventory_digest_input)
            assert d not in seen, f"digest 对某个分量不敏感: {v.regions}"
            seen.add(d)


# ═══════════════════════════════════════════════════════════════════════════
# §6 第 5 处阻塞：身份 binding 静态臂（Requirements 5.1 ~ 5.6 / Properties 4, 5）
# ═══════════════════════════════════════════════════════════════════════════


def build_binding(contract: Any, anchors: Any, *, entry_id: str = "probe/entry"):
    """直调 `_build_identity_binding`（绕开 DB —— 它只用 contract / anchors）。"""
    observer = OBS.PublishedIdentityObserver.__new__(OBS.PublishedIdentityObserver)

    class _Enum:
        value = "projection_contract"

    class _Bundle:
        bundle_id = "b"
        bundle_sha256 = "s"
        authority_model = _Enum()
        authority_model_definition_id = "amd"
        typed_slot_inventory: tuple = ()

    class _Res:
        representation_id = "r"
        content_version_id = "c"
        representation_generation = 1
        adapter_id = "a"
        adapter_build_digest = "d"
        bundle = _Bundle()

    return OBS.PublishedIdentityObserver._build_identity_binding(
        observer,
        contract=contract,
        anchors=anchors,
        dynamic_bindings={},
        entry_id=entry_id,
        resolution=_Res(),
        correlation_id="cid",
    )


class TestStaticIdentityBindingArm:
    """静态臂的正路径与两条显式失败（Requirements 5.1 ~ 5.4）。"""

    def test_static_anchor_yields_defined_name_binding(
        self, canary_contract: Any, canary_anchors: list[dict[str, str]]
    ) -> None:
        binding = build_binding(canary_contract, canary_anchors[0])
        assert binding.defined_name == canary_anchors[0]["defined_name"]
        assert binding.metadata_sheet == OBS.GT_SYNC_SHEET_NAME
        assert binding.table_key
        # `kind` 是派生属性：有 defined_name 即 static_region ⇒ 下游零改动的依据
        from app.services.workpaper_sync.excel_extract import BindingKind

        assert binding.kind == BindingKind.static_region

    def test_binding_matches_the_providers_own_static_bindings(
        self, a51: Any, canary_contract: Any, canary_anchors: list[dict[str, str]]
    ) -> None:
        """形态与 provider 现有 `static_identity_bindings()` 一致（Requirement 5.5 的依据）。"""
        by_name = {b.defined_name: b for b in a51.static_identity_bindings()}
        for anchor in canary_anchors:
            got = build_binding(canary_contract, anchor)
            expected = by_name[anchor["defined_name"]]
            assert (got.defined_name, got.table_key, got.metadata_sheet) == (
                expected.defined_name, expected.table_key, expected.metadata_sheet
            )

    def test_no_static_table_in_contract_is_frozen_child_unusable(
        self, canary_contract: Any, canary_anchors: list[dict[str, str]]
    ) -> None:
        """Requirement 5.3：契约无 `row_identity is None` 的静态表 ⇒ 显式失败。"""
        class NoStatic:
            contract_id = canary_contract.contract_id
            sheets: tuple = ()

        with pytest.raises(OBS.FrozenChildUnusableError) as exc:
            build_binding(NoStatic(), canary_anchors[0])
        assert "静态表" in str(exc.value)

    def test_ambiguous_alignment_fails_with_the_matched_set_in_the_message(
        self, canary_contract: Any, canary_anchors: list[dict[str, str]]
    ) -> None:
        """Requirement 5.4：对齐数 ≠ 1 ⇒ 显式失败且消息含匹配集合，**禁静默取首元素**。"""
        anchor = dict(canary_anchors[0])
        anchor["sheet_key"] = "不存在的-sheet-key"
        with pytest.raises(OBS.FrozenChildUnusableError) as exc:
            build_binding(canary_contract, anchor)
        msg = str(exc.value)
        assert "不存在的-sheet-key" in msg and "[]" in msg
        assert "不得随手挑第一张" in msg

    def test_static_arm_has_no_len_one_shortcut(self) -> None:
        """动态臂有 `elif len(row_tables) == 1` 的兜底；静态臂**刻意没有**。

        少了这条判据，有人「照抄动态臂」把兜底也抄过来，`len(matched) != 1` 就成了装饰。
        """
        import inspect

        src = inspect.getsource(OBS.PublishedIdentityObserver._build_identity_binding)
        static_seg = src[src.index('region_kind") or ""'):src.index("row_tables = [")]
        assert "len(matched) != 1" in static_seg
        assert "static_tables[0]" not in static_seg, "静态臂出现了「取首元素」兜底"
        assert "elif len(" not in static_seg


class TestDynamicBindingPathPreserved:
    """对照组：动态路径逐条原样保留（Requirement 5.2）。"""

    def test_dynamic_judgements_and_both_anchor_reads_are_intact(self) -> None:
        import inspect

        src = inspect.getsource(OBS.PublishedIdentityObserver._build_identity_binding)
        dyn = src[src.index("row_tables = ["):]
        assert "契约未声明任何带 row_identity 的表" in dyn
        assert 'anchors["table_name"]' in dyn
        assert 'anchors["uuid_column_letter"]' in dyn
        assert "elif len(row_tables) == 1:" in dyn
        assert "不得随手挑第一张" in dyn

    def test_static_arm_precedes_dynamic_path(self) -> None:
        import inspect

        src = inspect.getsource(OBS.PublishedIdentityObserver._build_identity_binding)
        assert src.index('region_kind") or ""') < src.index("row_tables = [")


class TestProperty5StaticBindingUniqueness:
    """Property 5: 静态主 binding 唯一性
    （Validates: Requirements 5.1, 5.3, 5.4）。"""

    @PBT
    @given(anchor_index=st.integers(min_value=0, max_value=1),
           sheet_key_mode=st.sampled_from(["aligned", "unknown", "duplicated"]))
    def test_either_uniquely_aligned_or_explicitly_failed(
        self, anchor_index: int, sheet_key_mode: str, canary_contract: Any,
        canary_anchors: list[dict[str, str]]
    ) -> None:
        from app.services.workpaper_sync.contracts import parse_contract

        anchor = dict(canary_anchors[anchor_index % len(canary_anchors)])
        contract = canary_contract
        if sheet_key_mode == "unknown":
            anchor["sheet_key"] = "no-such-key"
        elif sheet_key_mode == "duplicated":
            # 同一 sheet_key 上挂两张静态表 ⇒ 对齐数 2 ⇒ 必须显式失败
            payload = json.loads(json.dumps(dict(canary_contract.canonical_payload)))
            target = next(s for s in payload["sheets"]
                          if s["sheet_key"] == anchor["sheet_key"])
            clone = json.loads(json.dumps(target["tables"][0]))
            clone["table_key"] = f"{clone['table_key']}_dup"
            # stable_field_key 跨 sheet/table 必须唯一 ⇒ 克隆表的字段键一并改名，
            # 否则失败原因会变成 ContractSchemaError（那不是本属性要抓的东西）。
            for field in clone.get("fields") or []:
                field["stable_field_key"] = f"{field['stable_field_key']}--dup"
                field["json_pointer"] = f"{field['json_pointer']}--dup"
            target["tables"].append(clone)
            contract = parse_contract(payload, adapter_id=canary_contract.contract_id)

        try:
            binding = build_binding(contract, anchor)
        except OBS.FrozenChildUnusableError as exc:
            assert sheet_key_mode != "aligned", f"对齐形态不该失败: {exc}"
            return
        assert sheet_key_mode == "aligned", (
            f"{sheet_key_mode} 形态本该显式失败，却静默返回了 {binding!r}"
        )
        assert binding.defined_name == anchor["defined_name"]


class TestProperty4StaticAnchorRoundTrip:
    """Property 4: 静态锚点往返
    （Validates: Requirements 2.7, 5.6）。"""

    @PBT
    @given(region_count=st.integers(min_value=1, max_value=2))
    def test_declared_sheet_names_are_read_back_through_the_anchors(
        self, region_count: int, a51: Any, canary_instrumented: bytes,
        canary_contract: Any
    ) -> None:
        """注入后经 `_frozen_sheet_anchors` + `_collect_static_region_physical`
        读回的 `(sheet_key → 物理 sheet 名)` 映射等于 spec 声明。"""
        spec = a51.static_only_instrumentation_spec()
        declared = {r.sheet_key: r.excel_name for r in spec.static_regions}
        anchors = OBS._frozen_sheet_anchors(a51.build_instrumentation_payload())
        readback = {
            a["sheet_key"]: OBS._collect_static_region_physical(
                data=canary_instrumented, contract=canary_contract,
                key=a["sheet_key"], defined_name=a["defined_name"],
            )
            for a in anchors[:region_count]
        }
        assert readback == {k: v for k, v in declared.items() if k in readback}
        assert len(readback) == region_count


# ═══════════════════════════════════════════════════════════════════════════
# §7 寄生静态区与动态 entry 的零回归（Requirement 7 / Property 7）
# ═══════════════════════════════════════════════════════════════════════════

_XI_REL = "backend/app/services/workpaper_sync/excel_instrumentation.py"


def zip_entry_map(data: bytes) -> dict[str, bytes]:
    """zip 内每个部件的字节。

    🔴 比「原始 zip 字节」更正确的「逐字节相等」口径：`ZipFile.writestr(str_name, …)`
    把**当前时间**写进每条 ZipInfo 的 `date_time`，故同一逻辑跨秒运行两次原始字节就不等。
    那是打包元数据的噪声，不是 instrumentation 产物内容的差异。
    """
    with zipfile.ZipFile(io.BytesIO(data)) as zf:
        return {n: zf.read(n) for n in zf.namelist()}


@pytest.fixture(scope="module")
def head_instrumentation_module() -> Any:
    """把 git HEAD 版的 `excel_instrumentation.py` 作为独立模块加载 —— 真实「改造前」。"""
    import subprocess
    import types

    res = subprocess.run(
        ["git", "show", f"HEAD:{_XI_REL}"],
        cwd=_BACKEND.parent, capture_output=True, check=False,
    )
    if res.returncode != 0 or not res.stdout:
        pytest.skip(f"git show 取不到 HEAD 版 {_XI_REL}（未跟踪或无 git）")
    mod = types.ModuleType("_xi_head_for_zero_regression")
    mod.__file__ = str(_BACKEND / "app" / "services" / "workpaper_sync"
                       / "excel_instrumentation.py")
    sys.modules[mod.__name__] = mod
    exec(compile(res.stdout.decode("utf-8"), mod.__file__, "exec"), mod.__dict__)
    return mod


@pytest.fixture(scope="module")
def parasitic_entries() -> list[tuple[str, Any, list[Any]]]:
    """「动态 primary + 非空 `static_sheets`」entry 的**完整**分母（现算，禁写死成员）。

    发现面 = `workpaper_sync` 下源码含 `static_sheets=` 的模块；判据 = 该模块的
    `instrumentation_specs()` / `instrumentation_spec()` 里真有 spec 带非空 `static_sheets`。
    """
    import importlib

    svc = _BACKEND / "app" / "services" / "workpaper_sync"
    out: list[tuple[str, Any, list[Any]]] = []
    for py in sorted(svc.glob("*.py")):
        text = py.read_bytes().decode("utf-8", "replace")
        if not re.search(r"static_sheets\s*=", text):
            continue
        try:
            mod = importlib.import_module(
                f"app.services.workpaper_sync.{py.stem}"
            )
        except Exception:  # noqa: BLE001
            continue
        specs_fn = getattr(mod, "instrumentation_specs", None)
        spec_fn = getattr(mod, "instrumentation_spec", None)
        specs: list[Any] = []
        try:
            if callable(specs_fn):
                specs = list(specs_fn())
            elif callable(spec_fn):
                specs = [spec_fn()]
        except Exception:  # noqa: BLE001
            continue
        if any(getattr(s, "static_sheets", ()) for s in specs):
            out.append((str(getattr(specs[0], "entry_id", py.stem)), mod, specs))
    assert out, "寄生静态区 entry 分母为空 ⇒ Property 7 失去分母（须重算发现面）"
    return out


class TestProperty7ParasiticStaticRegionsZeroRegression:
    """Property 7: 寄生静态区零回归
    （Validates: Requirements 7.2, 7.3）。

    分母是 Task 1.1 现算的**完整**清单（由 `parasitic_entries` fixture 现算，禁写死成员）。
    """

    @PBT
    @given(pick=st.integers(min_value=0, max_value=32))
    def test_injection_artifact_is_part_for_part_identical_to_head(
        self, pick: int, parasitic_entries: list[tuple[str, Any, list[Any]]],
        head_instrumentation_module: Any, gate: Any
    ) -> None:
        import dataclasses as _dc

        entry_id, mod, specs = parasitic_entries[pick % len(parasitic_entries)]
        # 🔴 权威字节从 **spec 自己的** `template_relative_path` 取，不从 provider 的
        #    `read_authoritative_template()` —— 现算实证：寄生分母里的
        #    `phase5_d1_expansion` 是子模块，根本没有那个函数（entry 的读册入口在
        #    `phase5_d1_notes_receivable`）。按 provider 取会把判据卡在「模块没这函数」上。
        rel = specs[0].template_relative_path.replace("\\", "/")
        rel = rel.split("backend/wp_templates/")[-1].lstrip("/")
        source = (_BACKEND / "wp_templates" / rel).read_bytes()
        head = head_instrumentation_module

        head_specs = [
            head.ExcelInstrumentationSpec(**{
                f.name: getattr(s, f.name) for f in _dc.fields(s)
            })
            for s in specs
        ]
        head_art = head.instrument_workbook_bytes_multi(
            source, head_specs, gate=head.ExcelIdentityCarrierGate.load()
        )
        now_art = _xi.instrument_workbook_bytes_multi(source, specs, gate=gate)

        before, after = zip_entry_map(head_art.instrumented_bytes), zip_entry_map(
            now_art.instrumented_bytes
        )
        assert sorted(before) == sorted(after), (
            f"{entry_id}: 部件集合变了 "
            f"新增 {sorted(set(after) - set(before))} 丢失 {sorted(set(before) - set(after))}"
        )
        differing = [n for n in before if before[n] != after[n]]
        assert not differing, f"{entry_id}: 以下部件字节不等 {differing}"
        assert head_art.gt_sync_pairs == now_art.gt_sync_pairs
        assert head_art.defined_name_refs == now_art.defined_name_refs
        assert head_art.table_ref == now_art.table_ref
        assert dict(head_art.row_uuids) == dict(now_art.row_uuids)

    def test_denominator_covers_every_source_level_hit(
        self, parasitic_entries: list[tuple[str, Any, list[Any]]]
    ) -> None:
        """分母的**发现面**须与源码命中面对账：命中但不在分母的必须逐个有理由。

        现算已知的三类落差：引擎自身（`excel_instrumentation`）、本 spec 的 canary
        （纯静态、无动态 primary）、以及开关关闭的空分母成员。
        """
        svc = _BACKEND / "app" / "services" / "workpaper_sync"
        hits = {
            py.stem for py in svc.glob("*.py")
            if re.search(r"static_sheets\s*=",
                         py.read_bytes().decode("utf-8", "replace"))
        }
        in_denominator = {mod.__name__.rsplit(".", 1)[-1] for _e, mod, _s in parasitic_entries}
        unexplained = hits - in_denominator - {
            "excel_instrumentation",            # 引擎自身（字段定义 + 注入支）
            "phase5_a51_cashflow_audit",        # 本 spec 的 canary：纯静态，无动态 primary
            "phase5_e1_monetary_fund",          # 静态区受开关控制，现算空分母
        }
        assert not unexplained, (
            f"源码命中 `static_sheets=` 但既不在分母、也无登记理由: {sorted(unexplained)}"
        )

    def test_e1_static_region_is_an_empty_denominator_not_a_missed_member(self) -> None:
        """登记 design §1.9 的一处**过宽快照**：`E1` 被列进「寄生静态区成员」，
        但现算它的静态区受开关控制、现为空 ⇒ 它是**空分母**成员。

        本 spec 不动那个开关（不在范围内），只把事实登记下来，避免下一轮照抄过宽清单。
        """
        from app.services.workpaper_sync import phase5_e1_monetary_fund as e1

        assert e1.static_region_specs() == (), (
            "E1 的静态区已非空 ⇒ 它应进入 Property 7 的分母，本判据须随之更新"
        )
        # 附带登记：它的元素形态与注入支实际读的键**不同构**（开关打开会抛 KeyError）
        decl_src = __import__("inspect").getsource(e1._static_sheet_declarations)
        assert "region_boundary_locator" not in decl_src
        assert "excel_name" not in decl_src

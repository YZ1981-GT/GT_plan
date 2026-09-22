# -*- coding: utf-8 -*-
"""``workbook_read_scope`` 守卫：同一趟 extract 内 openpyxl 只解析一次，且句柄不泄漏。

背景（2026-09-22 真栈 cProfile 实测）：D4-营业收入的 ``ExcelSyncAdapter.extract`` 会
``for binding in _all_bindings()`` 跑 39 个 binding，每个 binding 读 ``data_only``
True/False 两个视图 ⇒ ``openpyxl.load_workbook`` 被调 **78 次**（累计 60.8s，其中
``apply_stylesheet`` 22.2s），是 store-projection 首请求十秒量级的主项。这 78 次读的是
**同一个文件的同一批字节**。

本文件把两条不变量钉住，二者缺一都会把优化变成 bug：

1. **少解析**：作用域内同一 ``(文件身份, data_only)`` 只 ``load_workbook`` 一次；
2. **不泄漏**：作用域退出时句柄必须全关 —— 否则 Windows 上 staged / repaired 临时文件
   随后的 ``unlink`` 会变成偶发 ``PermissionError``（性能优化变成数据通路故障）。

并且「只改算几次、不改算出什么」：缓存命中路径读出的格必须与无作用域路径**逐格恒等**。
"""
from __future__ import annotations

import openpyxl
import pytest

from app.services.workpaper_sync import excel_extract
from app.services.workpaper_sync.excel_extract import (
    _read_cell_view,
    workbook_read_scope,
)


@pytest.fixture()
def load_counter(monkeypatch: pytest.MonkeyPatch) -> list[tuple[str, bool]]:
    """记录每次真实 ``load_workbook`` 的 (文件名, data_only)。"""
    calls: list[tuple[str, bool]] = []
    real = openpyxl.load_workbook

    def counting(path, *args, **kwargs):  # type: ignore[no-untyped-def]
        calls.append((str(path), bool(kwargs.get("data_only"))))
        return real(path, *args, **kwargs)

    monkeypatch.setattr(openpyxl, "load_workbook", counting)
    return calls


def _write_book(path, *, sheet: str = "S1", a1: str = "hello", b2: float = 12.5):
    wb = openpyxl.Workbook()
    ws = wb.active
    ws.title = sheet
    ws["A1"] = a1
    ws["B2"] = b2
    wb.save(path)
    wb.close()
    return path


def _view(path, *, data_only: bool = True, sheet: str = "S1") -> dict:
    return _read_cell_view(
        path,
        sheet_name=sheet,
        columns=frozenset({"A", "B"}),
        rows=range(1, 3),
        data_only=data_only,
    )


# ═══ 1. 作用域内只解析一次 ═══


def test_scope_parses_the_same_workbook_view_only_once(tmp_path, load_counter) -> None:
    book = _write_book(tmp_path / "book.xlsx")

    with workbook_read_scope():
        first = _view(book)
        for _ in range(9):
            _view(book)

    assert len(load_counter) == 1, (
        f"作用域内同一 (文件, data_only) 解析了 {len(load_counter)} 次 —— "
        "workbook 复用未生效，78 次全簿解析的主项没被收掉"
    )
    assert first["A1"] == "hello" and first["B2"] == 12.5


def test_scope_keeps_data_only_views_separate(tmp_path, load_counter) -> None:
    """两个视图语义不同（计算值 vs 公式文本），必须各解析一次，绝不能互相命中。"""
    book = _write_book(tmp_path / "book.xlsx")

    with workbook_read_scope():
        _view(book, data_only=True)
        _view(book, data_only=True)
        _view(book, data_only=False)
        _view(book, data_only=False)

    assert sorted(flag for _, flag in load_counter) == [False, True], (
        f"data_only 两个视图的解析次数不对：{load_counter}"
    )


def test_scope_does_not_mix_up_distinct_files(tmp_path, load_counter) -> None:
    one = _write_book(tmp_path / "one.xlsx", a1="one")
    two = _write_book(tmp_path / "two.xlsx", a1="two")

    with workbook_read_scope():
        got_one = _view(one)
        got_two = _view(two)
        _view(one)
        _view(two)

    assert len(load_counter) == 2, f"不同文件应各解析一次：{load_counter}"
    assert got_one["A1"] == "one"
    assert got_two["A1"] == "two", "不同文件串用了缓存 —— 会读到另一个底稿的数据"


def test_cache_invalidates_when_the_file_is_rewritten(tmp_path, load_counter) -> None:
    """缓存键含 (mtime_ns, size)：文件被重写后必须重新解析，绝不返回旧内容。"""
    book = tmp_path / "book.xlsx"
    _write_book(book, a1="before")

    with workbook_read_scope():
        before = _view(book)
        _write_book(book, a1="after-rewritten-with-a-longer-value")
        after = _view(book)

    assert before["A1"] == "before"
    assert after["A1"] == "after-rewritten-with-a-longer-value", (
        "文件重写后仍读到旧字节 —— 缓存键没有跟住文件身份"
    )
    assert len(load_counter) == 2, f"重写后应重新解析：{load_counter}"


# ═══ 2. 句柄不泄漏 ═══


def test_scope_exit_releases_handles_so_the_artifact_can_be_deleted(tmp_path) -> None:
    """Windows 上句柄未关 ⇒ unlink 抛 PermissionError。staged/repaired 产物必须删得掉。"""
    book = _write_book(tmp_path / "book.xlsx")

    with workbook_read_scope():
        _view(book)
        _view(book, data_only=False)

    book.unlink()  # 句柄泄漏时此行在 Windows 上抛 PermissionError
    assert not book.exists()


def test_scope_exit_releases_handles_even_when_the_body_raises(tmp_path) -> None:
    book = _write_book(tmp_path / "book.xlsx")

    with pytest.raises(ZeroDivisionError):
        with workbook_read_scope():
            _view(book)
            raise ZeroDivisionError("模拟作用域内失败")

    book.unlink()
    assert not book.exists()


def test_scope_state_is_cleared_on_exit(tmp_path) -> None:
    book = _write_book(tmp_path / "book.xlsx")

    assert excel_extract._workbook_scope.get() is None
    with workbook_read_scope():
        _view(book)
        assert excel_extract._workbook_scope.get() is not None
    assert excel_extract._workbook_scope.get() is None, (
        "作用域退出后 ContextVar 未复位 —— 后续请求会复用上一次的句柄字典"
    )


def test_nested_scope_reuses_the_outer_cache_and_defers_closing(
    tmp_path, load_counter
) -> None:
    """嵌套进入不得提前关闭外层句柄（否则内层退出后外层继续读会炸）。"""
    book = _write_book(tmp_path / "book.xlsx")

    with workbook_read_scope():
        _view(book)
        with workbook_read_scope():
            _view(book)
        after_inner = _view(book)  # 内层退出后仍必须可读

    assert len(load_counter) == 1, f"嵌套作用域重复解析了：{load_counter}"
    assert after_inner["A1"] == "hello"
    book.unlink()


# ═══ 3. 作用域外行为不变 ═══


def test_without_a_scope_every_call_parses_and_closes(tmp_path, load_counter) -> None:
    book = _write_book(tmp_path / "book.xlsx")

    _view(book)
    _view(book)
    _view(book)

    assert len(load_counter) == 3, (
        f"作用域外不得跨调用共享句柄（会把打开的 zip 留到调用栈之外）：{load_counter}"
    )
    book.unlink()


# ═══ 4. 只改「算几次」不改「算出什么」 ═══


def test_scoped_and_unscoped_reads_are_cell_for_cell_identical(tmp_path) -> None:
    book = _write_book(tmp_path / "book.xlsx")

    unscoped_values = _view(book, data_only=True)
    unscoped_formulas = _view(book, data_only=False)
    with workbook_read_scope():
        scoped_values = _view(book, data_only=True)
        scoped_formulas = _view(book, data_only=False)
        scoped_values_again = _view(book, data_only=True)

    assert scoped_values == unscoped_values
    assert scoped_formulas == unscoped_formulas
    assert scoped_values_again == unscoped_values, "同一作用域内重复读出现漂移"


def test_missing_sheet_still_raises_the_narrow_identity_error(tmp_path) -> None:
    """作用域不得把窄类型错误吞成别的（sheet 不存在仍须是 IdentityCarrierMissingError）。"""
    book = _write_book(tmp_path / "book.xlsx", sheet="Real")

    with workbook_read_scope():
        with pytest.raises(excel_extract.IdentityCarrierMissingError):
            _view(book, sheet="Nope")
        # 同一作用域内后续正常读仍要可用（异常不得污染缓存）
        assert _view(book, sheet="Real")["A1"] == "hello"
    book.unlink()

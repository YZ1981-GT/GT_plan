# -*- coding: utf-8 -*-
"""🔴 D1-4 三区共存的硬取舍：静态区进契约 ⇒ 动态区不能再插行。

spec: d1-sync-row-table-engine-and-d1-coverage · Task 26 子目标 ①「同 sheet 位移链实证」

═══ 本文件钉的是一个**我自己引入的能力回退** ═══

O 节把 `bad_debt_notetype_rows`（D1-4 第三区，绝对行 23/24）加进契约，换来静态区 12 格可写。
代价当时没被发现：该区在 **footer（R22）之下**，于是两个动态区（individual R13-16 /
portfolio R18-21）**一旦需要插行就被引擎拒绝** ——
`RowSetDivergenceError[contract_static_row_below_insertion]`。

**为什么 O/P 两节的 14 条真栈判据看不到**：它们只喂第三区的载荷，从未构造「动态区行数
超过模板容量」的场景，于是 `row_shift is None` 这条断言是在**本来就不需要位移**的场景下成立的
—— 典型的「结构性零不配变异证明」（本 spec 方法论铁律 ⑮，我自己写的那条）。

═══ 平台对这类危险的态度是明确的，且 D4 已有相反先例 ═══

* `excel_materialize.py` 模块 docstring 把它登记为不可安全插行的第 (e) 类，
  并注明「**本 spec 实施中实测发现，design.md 未覆盖**」；
* 报错自带解除条件：「extract 侧的静态行定位同样变成位移感知（读写两侧一起改）」；
* 🔴 **`phase5_d4_policy_check_sheet.py` 撞过同一问题并选了相反处置** ——
  「信用/说明/结论在 footer 之下：引擎对 `contract_static_row_below_insertion` fail-closed
  （extract 仍按死行号反读），故**不入契约**；仍由 HTML 持久化。待 marker-relative content
  字段落地后再扩。」

⇒ D1-4 与 D4 在**结构相同**的情形上做了**相反**的选择。本文件不替产品裁决，
但把两边的事实都钉住，让这条不一致不能被忽略。

═══ 判据形态 ═══

对照/实验成对（这是本文件的核心，单边都不成立）：
* **对照**：把静态 table 从契约里滤掉 ⇒ 插行计划**能**算出来，且 `insert_at/count/style_from` 逐值确定；
* **实验**：完整契约 ⇒ 抛 `RowSetDivergenceError` 且消息里带 `contract_static_row_below_insertion`。

🔴 对照契约用「从 `build_contract_payload()` 过滤掉该 table」构造，**不用 `git show HEAD:`**
—— HEAD 会移动，判据会在别人提交后语义漂移。
"""
from __future__ import annotations

import copy
import hashlib
import io
import json
import sys
import zipfile
from functools import lru_cache
from pathlib import Path
from typing import Any

import openpyxl
import pytest

_BACKEND = Path(__file__).resolve().parents[2]
if str(_BACKEND) not in sys.path:
    sys.path.insert(0, str(_BACKEND))

from app.services.workpaper_sync import excel_instrumentation as EI  # noqa: E402
from app.services.workpaper_sync import excel_materialize as M  # noqa: E402
from app.services.workpaper_sync import phase5_d1_04_bad_debt as D104  # noqa: E402
from app.services.workpaper_sync import phase5_d1_expansion as EXP  # noqa: E402
from app.services.workpaper_sync import phase5_d1_notes_receivable as P  # noqa: E402
from app.services.workpaper_sync.contracts import parse_contract  # noqa: E402
from app.services.workpaper_sync.excel_extract import (  # noqa: E402
    ExcelIdentityBinding,
    _scan_row_identities,
    read_runtime_binding_pairs,
)

# 🔴 T7 裁决 A：静态 table 已撤回（`_INCLUDE_D104_NOTETYPE_STATIC=False`）。
#    本文件证明的是「静态 table **在**契约里 ⇒ 动态区插行被 fail-closed」这个**回退**，
#    撤回后该回退不再存在 ⇒ 整体跳过，**不删**（它是撤回决策的证据，开关翻回 True 时复活）。
#    撤回后的**镜像判据**在 `test_d104_static_region_excluded.py::
#    test_dynamic_region_can_insert_row_again_after_withdrawal` —— 那条证明「不拒」，
#    两条一起把「回退由静态 table 引入」的因果钉死在两个方向。
pytestmark = pytest.mark.skipif(
    not EXP._INCLUDE_D104_NOTETYPE_STATIC,
    reason="T7 裁决 A：静态 table 已撤回，本文件证明的回退不再存在；镜像判据见 test_d104_static_region_excluded.py",
)

SHEET = D104.MANAGED_SHEET_D104
STATIC_TABLE_KEY = "bad_debt_notetype_rows"

#: 模板几何（字面量，不从 spec 读 —— 守卫期望值必须独立于被守对象）
INDIVIDUAL_ROWS = (13, 16)   # 4 行容量
PORTFOLIO_ROWS = (18, 21)    # 4 行容量
FOOTER_ROW = 22
STATIC_ROWS = (23, 24)
#: 插入点 = individual 区末行 + 1
EXPECTED_INSERT_AT = 17
EXPECTED_STYLE_FROM = 16
#: 静态区贡献的受管格数（2 行 × 9 列）
STATIC_CELL_COUNT = 18


@lru_cache(maxsize=1)
def _substrate() -> bytes:
    return EI.instrument_workbook_bytes_multi(
        P.read_authoritative_template(),
        EXP.instrumentation_specs(),
        gate=P.excel_carrier_gate(),
    ).instrumented_bytes


def _payload_without_static() -> dict[str, Any]:
    """完整契约 payload 去掉静态 table —— 即「O 节之前」的契约形态。"""
    payload = copy.deepcopy(P.build_contract_payload())
    dropped = 0
    for sheet in payload.get("sheets") or []:
        tables = sheet.get("tables") or []
        keep = [t for t in tables if t.get("table_key") != STATIC_TABLE_KEY]
        dropped += len(tables) - len(keep)
        sheet["tables"] = keep
    assert dropped == 1, f"应恰好滤掉 1 张静态 table，实际 {dropped}"
    return payload


def _minted_ids(spec) -> list[str]:
    ws = openpyxl.load_workbook(io.BytesIO(_substrate()), data_only=False)[SHEET]
    return [
        str(ws[f"{spec.uuid_col}{r}"].value)
        for r in range(spec.first_data_row, spec.last_data_row + 1)
        if ws[f"{spec.uuid_col}{r}"].value
    ]


def _scan(spec, *, contract):
    sub = _substrate()
    ws = openpyxl.load_workbook(io.BytesIO(sub), data_only=False)[SHEET]
    raw_by_row = {
        r: str(ws[f"{spec.uuid_col}{r}"].value or "")
        for r in range(spec.first_data_row, spec.last_data_row + 1)
    }
    table = next(t for s in contract.sheets for t in s.tables if t.table_key == spec.table_key)
    return _scan_row_identities(
        table=table, sheet_name=SHEET, uuid_column=spec.uuid_col,
        raw_by_row=raw_by_row, artifact_sha256=hashlib.sha256(sub).hexdigest(),
        tombstoned=(),
    )


def _notetype_payload() -> str:
    return json.dumps(
        [
            {"rowId": row_id, "noteType": label, "isFixed": True,
             "priorUnadjusted": 1000.0 * i, "priorAje": 11.0 * i, "priorRje": 22.0 * i,
             "currentUnadjusted": 2000.0 * i, "currentAje": 33.0 * i, "currentRje": 44.0 * i}
            for i, (row_id, _no, label) in enumerate(D104.NOTETYPE_FIXED_ROWS, start=1)
        ],
        ensure_ascii=False,
    )


def _plan(*, payload: dict[str, Any], n_extra: int, with_static: bool = False):
    """按给定契约 payload 规划 individual 区写入；n_extra = 超出模板容量的新行数。

    `with_static=True` 时同时喂第三区载荷 —— 只有喂了投影才有值，
    否则 `_emit` 对缺键 `return`，静态格不会出现在写入面（这是**设计**，不是缺陷）。
    """
    sub = _substrate()
    contract = parse_contract(payload, adapter_id=P.ADAPTER_ID)
    spec = D104.SPEC_D104_INDIVIDUAL
    ids = _minted_ids(spec) + [f"GTROW-NEW-{i:04d}" for i in range(1, n_extra + 1)]
    rows = [
        {"rowId": rid, "label": f"客户{i}", "priorUnadjusted": 100.0 + i,
         "currentProvision": 10.0 + i}
        for i, rid in enumerate(ids, start=1)
    ]
    payloads: dict[str, str] = {spec.store_item_id: json.dumps(rows, ensure_ascii=False)}
    if with_static:
        payloads[D104.SPEC_D104_NOTETYPE.store_item_id] = _notetype_payload()
    proj = P.build_combined_store_projection(payloads, contract=contract)
    binding = ExcelIdentityBinding(
        table_key=spec.table_key, table_name=spec.table_name,
        uuid_column=spec.uuid_col, metadata_sheet="_GT_SYNC",
    )
    with zipfile.ZipFile(io.BytesIO(sub)) as zf:
        region = M.resolve_managed_region(zf, contract=contract, binding=binding)
    return M.plan_managed_writes(
        projection=proj, contract=contract, binding=binding, region=region,
        scan=_scan(spec, contract=contract),
        substrate_entries=M._read_entries(sub),
        substrate_formulas={},
        runtime_binding=read_runtime_binding_pairs(zipfile.ZipFile(io.BytesIO(sub))),
    )


# ═══════════════════════════════════════════════════════════════════════
# 1. 前提：模板几何真的是「静态区在 footer 之下」
# ═══════════════════════════════════════════════════════════════════════


def test_premise_static_region_sits_below_footer() -> None:
    """这是整条危险的成因 —— 不成立则本文件其余断言无从谈起。"""
    ind = D104.SPEC_D104_INDIVIDUAL
    por = D104.SPEC_D104_PORTFOLIO
    note = D104.SPEC_D104_NOTETYPE
    assert (ind.first_data_row, ind.last_data_row) == INDIVIDUAL_ROWS
    assert (por.first_data_row, por.last_data_row) == PORTFOLIO_ROWS
    assert ind.footer_row == por.footer_row == FOOTER_ROW
    assert (note.first_data_row, note.last_data_row) == STATIC_ROWS
    # 🔴 要害：静态区行号 > footer 行号
    assert note.first_data_row > FOOTER_ROW, (
        "静态区不再位于 footer 之下 ⇒ 本文件登记的取舍可能已消失，请复核"
    )
    # 静态区无行身份键（正是它必须走静态路径的原因）
    assert note.row_identity_key == ""


def test_premise_template_capacity_is_four_plus_four() -> None:
    """容量 = 4 + 4。前端无上限（见下），所以这个数就是暴露面的分母。"""
    assert INDIVIDUAL_ROWS[1] - INDIVIDUAL_ROWS[0] + 1 == 4
    assert PORTFOLIO_ROWS[1] - PORTFOLIO_ROWS[0] + 1 == 4
    assert len(_minted_ids(D104.SPEC_D104_INDIVIDUAL)) == 4, (
        "instrumented 模板的 individual 区不是 4 行 minted UUID"
    )


# ═══════════════════════════════════════════════════════════════════════
# 2. 对照 / 实验成对 —— 这是本文件的核心
# ═══════════════════════════════════════════════════════════════════════


def test_control_without_static_table_row_insertion_is_planned() -> None:
    """**对照**：静态 table 不在契约里 ⇒ 插行计划算得出来，且逐值确定。

    没有这一条，下面那条「加了静态 table 就拒绝」只能证明「拒绝了」，
    证明不了「**是静态 table 导致的**」—— 也许模板本来就不支持插行。
    """
    plan = _plan(payload=_payload_without_static(), n_extra=1)
    shift = plan.row_shift
    assert shift is not None, "对照组没算出插行计划 ⇒ 归因前提不成立"
    assert shift.insert_at == EXPECTED_INSERT_AT
    assert shift.count == 1
    assert shift.style_from == EXPECTED_STYLE_FROM
    assert shift.table_key == D104.SPEC_D104_INDIVIDUAL.table_key


def test_experiment_with_static_table_row_insertion_is_refused() -> None:
    """**实验**：完整契约（含静态 table）⇒ fail-closed，且必须是**那一类**原因。

    只断言「抛了 RowSetDivergenceError」不够 —— 引擎有五类不可安全插行情形，
    各自在消息里带自己的 `error_code` 片段。合并两类的改动必须能被这条抓住。
    """
    from app.services.workpaper_sync.excel_materialize import RowSetDivergenceError

    with pytest.raises(RowSetDivergenceError) as ei:
        _plan(payload=P.build_contract_payload(), n_extra=1)
    msg = str(ei.value)
    assert "contract_static_row_below_insertion" in msg, (
        f"拒绝了但不是静态行那一类，归因失效：{msg[:200]}"
    )
    # 消息必须点名真实的插入点与静态格数量（否则读者无从判断影响面）
    assert str(EXPECTED_INSERT_AT) in msg
    assert f"共 {STATIC_CELL_COUNT} 个" in msg, (
        f"静态格数量不是 {STATIC_CELL_COUNT} ⇒ 静态区形态变了，请复核：{msg[:200]}"
    )
    # 平台自带的解除条件必须还在（它是后续裁决的依据，不能悄悄消失）
    assert "位移感知" in msg


def test_refusal_is_triggered_by_capacity_not_by_new_row_ids() -> None:
    """不超容量时（哪怕行身份全是模板已 mint 的）**不**触发拒绝。

    🔴 这条把「任何新 rowId 都会被拒」与「超容量才被拒」分开 ——
    我的第一版探针用合成 rowId 填满 4 行就打红，误以为「连满容量都不行」，
    实际是因为合成 id 与 minted UUID 不匹配、4 行全被当新行。
    """
    plan = _plan(payload=P.build_contract_payload(), n_extra=0, with_static=True)
    assert plan.row_shift is None, "不超容量却算出了插行计划"
    # 静态区的 12 格（6 可编辑列 × 2 行）与动态区写入**同处一个 plan**
    coords = {w.coord for w in plan.writes}
    static_coords = sorted(
        c for c in coords if c[1:].isdigit() and int(c[1:]) in STATIC_ROWS
    )
    assert len(static_coords) == 12, f"静态区写入面应为 12 格，实际 {static_coords}"
    assert {"A23", "A24", "L23", "M24"} <= coords
    # 动态区也真的写了（否则「同处一个 plan」这句没意义）
    assert {"A13", "A16"} <= coords


# ═══════════════════════════════════════════════════════════════════════
# 3. 暴露面：前端无容量上限 + D4 的相反先例
# ═══════════════════════════════════════════════════════════════════════

_FE = _BACKEND.parent / "audit-platform/frontend/src/components/workpaper"


def _strip_comments(src: str) -> str:
    import re

    return (
        re.sub(r"(^|[^:])//[^\n]*", r"\1",
               re.sub(r"/\*[\s\S]*?\*/", "", re.sub(r"<!--[\s\S]*?-->", "", src)))
    )


def test_frontend_addsubrow_has_no_capacity_cap() -> None:
    """暴露面量化：前端 `addSubRow` 只挡 readonly，**没有**容量上限。

    ⇒ 审计师加第 5 个「按单项子行」就会让本 entry 在切「在线编辑」时 materialize 失败。
    失败是可见的（不是静默写坏数据），但发生在**切换时**而不是**加行时**。
    🔴 本条只登记事实、不替产品裁决是否要加上限（那会移除合法业务能力：
       个别认定的客户数天然可以超过 4 个）。
    """
    src = _strip_comments((_FE / "composables/useD1BadDebt.ts").read_text(encoding="utf-8"))
    idx = src.index("function addSubRow")
    body = src[idx:idx + 1200]
    assert "isReadonly" in body, "addSubRow 连 readonly 都不挡了？请复核"
    for token in ("MAX_", "capacity", "length >=", "length >", "TEMPLATE_ROWS"):
        assert token not in body, (
            f"addSubRow 里出现了 {token!r} ⇒ 可能已加容量上限，本条登记的暴露面需更新"
        )
    # 正面对照：确认扫的确实是「无界追加」那一句
    assert "individualRows.value = [...individualRows.value, newRow]" in body


def test_d4_precedent_chose_the_opposite_and_is_still_recorded() -> None:
    """🔴 D4 在**结构相同**的情形上选了相反处置，该记录必须还在。

    D1-4 把 footer 之下的静态区**放进**契约（换来 12 格可写、代价是动态区不能插行）；
    D4 的 policy check sheet 把同类区**排除在**契约外（保住插行、代价是那些格 HTML-only）。
    两边不一致本身不是缺陷，但必须**可见** —— 后续裁决要在这两条之间选。
    本条防的是「D4 那段说明被删掉后，这条不一致变成无人知晓的隐性差异」。
    """
    src = (_BACKEND / "app/services/workpaper_sync/phase5_d4_policy_check_sheet.py").read_text(
        encoding="utf-8"
    )
    assert "contract_static_row_below_insertion" in src, (
        "D4 policy check sheet 不再提及该 error_code ⇒ 先例记录可能已被删，"
        "D1-4 与 D4 的处置不一致会失去唯一书面依据"
    )
    assert "不入契约" in src, "D4 的「不入契约」裁决措辞已变，请复核两边一致性"


def test_platform_documents_this_hazard_as_a_named_category() -> None:
    """引擎模块 docstring 必须仍把它登记为具名的一类（不是匿名 fail-closed）。

    具名是可归因的前提：五类不可安全插行情形两两可分辨，
    合并成一个笼统错误会让 `test_experiment_...` 的归因断言失效。
    """
    src = (_BACKEND / "app/services/workpaper_sync/excel_materialize.py").read_text(
        encoding="utf-8"
    )
    assert "contract_static_row_below_insertion" in src
    # 解除条件必须写在代码里（本文件与 tasks 登记都引用它）
    assert "位移感知" in src

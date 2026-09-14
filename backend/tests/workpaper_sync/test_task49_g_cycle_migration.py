# -*- coding: utf-8 -*-
r"""
test_task49_g_cycle_migration — G 循环（除 G7）Excel 独立 entry 迁移验证

spec: workpaper-html-onlyoffice-bidirectional-writeback-closure · Wave 5 Task 49
Requirements: 6.1, 6.2, 6.10, 9.1, 12.1, 12.4, 12.10, 12.11, 12.12, 14.1
             （另覆盖 1.4 / 1.7 / 12.8 / 12.9）

验证 Properties:
  - Property 20: generated col 占位不可注册生产 adapter
    （本 slice 17 条 entry 的 contract 数为 0 ⇒ **不宣称通过**；正向核验落在 G 循环唯一
      的生产契约 g7.soe_subsidiary_disclosure.json 上，它属**已排除**的 G7）
  - Property 21: contract 字段完整（同上，**不宣称通过**，见 TestProperty20And21）
  - Property 28: immutable definition 漂移 fail closed（真分母：15 个模板 digest +
    13 个 normalized_structure_hash + 17 组 sheet!cell 真读 + G1 的 18 个 tab 名比对）
  - Property 69: evidence 由逐 scenario 实体与服务端重算闭合（负向分母真验，
    正向逐 scenario 闭合不宣称通过，已登记 BP-4）
  - Property 70: scenario/evidence/contract/bundle 不得跨 entry 复用

═══ 判据设计（避开 Task 46 首轮踩过的四类坑，并修正 Task 48 对 G 循环不成立的假设）═══

1. **不用空分母重言式**。「G 循环没有 contract 所以 Property 21 通过」是重言式。本文件对
   每条 Property 显式区分「有真分母的部分（真验）」与「无适用分母的部分（只断言前提成立 +
   承载者存在，**不宣称通过**）」，划分与 slice 的 `property_denominators` 逐项对齐，
   类名与 docstring 明写不宣称的部分。

2. **不用 fail-open**。没有 `pytest.skip`、没有 `if x: assert ...`（缺值即跳过）、
   没有 `except Exception` 吞异常。缺文件/缺字段一律打红。

3. **不用硬编码豁免**。本 slice 没有 pilot，也不给任何 entry 开例外列。

4. **不假设与 F 循环同形**。G 循环实测出五处与 F/D 不同的形态，判据逐 entry 从 slice 读
   声明再回源码核对，而不是写死一个值（详见 slice 的 `g_cycle_form_differences`）：
   * payload 列有四种（remark 单写 / conclusion 单写 / storage-contract 双写 / 写入点双写）
     —— 写死 `payload_column == "remark"`（F 的做法）会让 10 条 entry 判错；
   * HTTP 客户端两族（`api` from apiProxy 13 条 / `http` from utils/http 4 条）
     —— F 的探针 `api\.get\(` 对那 4 条必然失配；
   * 一本工作簿服务多条 entry（G4/G6 各 3 条）—— F 的 belongs_to_entry 单射断言不成立；
   * 行身份三族（生成前缀串 / crypto.randomUUID / 源模板固定 rowKey）
     —— F 写死 `row_identity_key in ('rowId','id')` 会把 G14 的 `rowKey` 判非法；
   * 传输键声明散落在多模块（80 个重复字面量）—— 需要可复算的规模判据而不是逐条豁免。

依赖（缺任一即 fail closed）：
  - backend/data/workpaper_sync_g_cycle_manifest_slice.json（frozen slice）
  - backend/data/workpaper_sync_g_cycle_deletion_plan.json（deletion plan）
  - backend/data/workpaper_sync_migration_paradigm.json（Task 45 范式，只读）
  - backend/data/workpaper_sync_entry_manifest.json（source-backed manifest，只读）
  - backend/wp_templates/G/ 与 backend/wp_templates/_index.json（运行时权威，唯一真源）
"""
from __future__ import annotations

import hashlib
import json
import pathlib
import re
import sys
from typing import Any, Callable

import pytest

# ────────────────────────────────────────────────────────────────────────────
# Paths
# ────────────────────────────────────────────────────────────────────────────
# __file__ = backend/tests/workpaper_sync/test_...py
# parents[0] = workpaper_sync/, [1] = tests/, [2] = backend/, [3] = repo root
_THIS = pathlib.Path(__file__).resolve()
ROOT = _THIS.parents[3]
BACKEND = ROOT / "backend"
FRONTEND = ROOT / "audit-platform" / "frontend" / "src"
WP_COMPONENTS = FRONTEND / "components" / "workpaper"
COMPOSABLES = WP_COMPONENTS / "composables"
GLOBAL_COMPOSABLES = FRONTEND / "composables"
SYNC_DIR = WP_COMPONENTS / "sync"
DATA = BACKEND / "data"

MANIFEST_SLICE_PATH = DATA / "workpaper_sync_g_cycle_manifest_slice.json"
DELETION_PLAN_PATH = DATA / "workpaper_sync_g_cycle_deletion_plan.json"
PARADIGM_PATH = DATA / "workpaper_sync_migration_paradigm.json"
FULL_MANIFEST_PATH = DATA / "workpaper_sync_entry_manifest.json"
D_SLICE_PATH = DATA / "workpaper_sync_d_cycle_manifest_slice.json"
E_SLICE_PATH = DATA / "workpaper_sync_e_cycle_manifest_slice.json"
F_SLICE_PATH = DATA / "workpaper_sync_f_cycle_manifest_slice.json"
CONTRACT_DIR = DATA / "workpaper_sync_contracts"
G7_PILOT_CONTRACT = CONTRACT_DIR / "g7.soe_subsidiary_disclosure.json"
TEMPLATE_DIR = BACKEND / "wp_templates"
TEMPLATE_INDEX = TEMPLATE_DIR / "_index.json"
CHECKLIST_ROUTER = BACKEND / "app" / "routers" / "checklist_responses.py"
TEMPLATE_FINDER = BACKEND / "app" / "services" / "wp_template_finder.py"
REGISTRY = BACKEND / "app" / "services" / "workpaper_sync" / "adapters" / "registry.py"
HTML_RENDERER_REGISTRY = WP_COMPONENTS / "htmlRendererRegistry.ts"
G1_SHEET_LABELS = COMPOSABLES / "g1SheetLabels.ts"
G1_DUAL_MODE = COMPOSABLES / "useG1DualMode.ts"
E1_HOST = WP_COMPONENTS / "GtE1MonetaryFund.vue"
G6_NOTE_SECTION_MAP = COMPOSABLES / "g6NoteSectionMap.ts"
G4_NOTE_SECTION_MAP = COMPOSABLES / "g4NoteSectionMap.ts"
G6_LISTED_TAB = (
    WP_COMPONENTS / "g6-other-bond-investment-main" / "core" / "G6TabDisclosureListed.vue"
)
G6_LISTED_ROWS = COMPOSABLES / "g6ListedDisclosureRows.ts"
G6_LISTED_COMPOSABLE = COMPOSABLES / "useG6DisclosureListed.ts"
NOTE_SYNC_REGISTRY = DATA / "note_workpaper_sync_registry.json"
UNREACHABLE_STUB = WP_COMPONENTS / "GtG6OtherBondEcl.vue"
SHARED_BASE = COMPOSABLES / "useWorkpaperEntryDualMode.ts"

#: AC 1.3 的能力态枚举（真源在范式 JSON 的 adjudication_criteria.capability_enum）。
CAPABILITY_ENUM = ("bidirectional", "single_html", "single_onlyoffice", "unreachable")
#: step 3 的二值结论（真源在范式 JSON 的 anti_patterns[AP-1].allowed_verdict_values）。
HTML_COUNTERPART_VERDICTS = ("none", "exists")

#: AC 1.4 的 UI 义务落点（Task 46 收口新建的单一真源，本任务复用不新建第二份）。
NOTICE_MODULE = SYNC_DIR / "workpaperEntrySyncNotice.ts"
NOTICE_COMPONENT = SYNC_DIR / "GtEntrySyncCapabilityNotice.vue"
NOTICE_COMPONENT_NAME = "GtEntrySyncCapabilityNotice"

#: G7 的三条 entry —— 任务正文明令排除，本文件的多处前提依赖「它们不在 slice 里」。
G7_ENTRY_IDS = frozenset(
    {
        "xlsx/gt-g7-long-term-equity-main",
        "xlsx/gt-g7-equity-method",
        "xlsx/gt-g7-equity-subsidiary",
    }
)
#: 不可达旧桩（AC 1.7）——它的 independent_entry=false，不进 slice。
UNREACHABLE_STUB_ENTRY_ID = "xlsx/gt-g6-other-bond-ecl"

#: 附注不得同步的 metadata sheet。
FORBIDDEN_NOTE_SHEETS = ("底稿目录", "GT_Custom", "修订说明")


# ────────────────────────────────────────────────────────────────────────────
# Helpers
# ────────────────────────────────────────────────────────────────────────────
def _load(path: pathlib.Path) -> dict:
    return json.loads(path.read_text(encoding="utf-8"))


def _sha256_of(path: pathlib.Path) -> str:
    return hashlib.sha256(path.read_bytes()).hexdigest()


def _strip_ts_comments(source: str) -> str:
    """剥掉 TS/Vue 注释 —— 说明文字不得充当判据证据。

    🔴 剥完必须反向自检（见 TestGuardSelfChecks），剥过头会让判据恒真。本文件对它的
    依赖比 Task 48 更重：G6 的自造形态判据**正是**「剥注释后必须为空」，剥不干净会假绿、
    剥过头会假红。
    """
    source = re.sub(r"/\*[\s\S]*?\*/", "", source)
    source = re.sub(r"(?m)^\s*//.*$", "", source)
    source = re.sub(r"(?m)//[^\n\"'`]*$", "", source)
    source = re.sub(r"<!--[\s\S]*?-->", "", source)
    return source


def _vue_template(source: str) -> str:
    """取 SFC 的**外层** template 区块。

    🔴 不能用 `source.split("</template>")[0]`：SFC 里嵌套的具名插槽/`<template v-else>`
    会先闭合，第一处 `</template>` 落在内层 ⇒ 截出来的片段比真实模板短，判据会把
    「已在模板里」误判成「不在模板里」。这里改用 `<script` 边界。
    """
    marker = source.find("<script")
    assert marker > 0, "找不到 <script 边界 —— 该文件不是常规 SFC（template 在 script 之前）"
    head = source[:marker]
    assert "<template>" in head, "SFC 头部没有 <template>"
    return head


def _toolbar_block(template: str, gate_anchor: str) -> str:
    """按 slice 声明的工具栏门控锚点截出该区块（开标签行 → 同缩进的 `</div>`）。

    🔴 为什么按 slice 声明取锚点而不是写死正则：G 循环 17 个宿主的工具栏 class 逐个不同
    （g1-toolbar / g2-interest-receivable-toolbar / … / g14-credit-impairment-loss-toolbar），
    其中 6 个还带 `v-if="currentSheet !== '底稿目录'"` 前缀、4 个宿主另有一个 `gN-index-toolbar`
    兄弟区块。写死任一种都会让其余的判据静默恒真（挂载点根本不在被搜的区块里也不报错）。
    """
    lines = template.splitlines()
    opens = [i for i, ln in enumerate(lines) if gate_anchor in ln]
    assert len(opens) == 1, (
        f"工具栏门控锚点 {gate_anchor!r} 在模板里命中 {len(opens)} 次（应为 1）"
    )
    start = opens[0]
    indent = len(lines[start]) - len(lines[start].lstrip())
    for i in range(start + 1, len(lines)):
        if lines[i].strip() == "</div>" and (len(lines[i]) - len(lines[i].lstrip())) == indent:
            return "\n".join(lines[start : i + 1])
    pytest.fail(f"锚点 {gate_anchor!r} 之后找不到同缩进的 </div> —— 工具栏区块无法界定")


def _cell_value(path: pathlib.Path, sheet: str, cell: str) -> Any:
    """openpyxl 真读权威模板的一个单元格（data_only）。"""
    import openpyxl
    from openpyxl.utils import column_index_from_string

    workbook = openpyxl.load_workbook(path, data_only=True)
    try:
        assert sheet in workbook.sheetnames, (
            f"{path.name} 里没有 sheet {sheet!r}；实有 {workbook.sheetnames}"
        )
        match = re.fullmatch(r"([A-Z]+)(\d+)", cell)
        assert match, f"非法单元格坐标 {cell!r}"
        return workbook[sheet].cell(
            row=int(match.group(2)),
            column=column_index_from_string(match.group(1)),
        ).value
    finally:
        workbook.close()


def _sheet_names(path: pathlib.Path) -> list[str]:
    import openpyxl

    workbook = openpyxl.load_workbook(path, read_only=True)
    try:
        return list(workbook.sheetnames)
    finally:
        workbook.close()


def _indexed_relpaths() -> set[str]:
    """`_index.json` 登记的相对路径集合（分隔符归一为 /）。"""
    index = _load(TEMPLATE_INDEX)
    return {str(f["relative_path"]).replace("\\", "/") for f in index["files"]}


def _resolve_frontend(rel: str) -> pathlib.Path:
    """把 slice 里登记的仓库相对路径解析成绝对路径。"""
    return ROOT / rel.split("#")[0]


def _parse_string_record(body: str) -> list[tuple[str, str]]:
    """把 TS 对象字面量的 `key: 'value'` 逐对解析出来（值允许换行到下一行）。

    🔴 为什么不用 `re.findall(r"'([^']{20,})'")` 之类的「捞长字符串」：那把**长度**当
    「有实质内容」的代理，短理由（`'同上（续表）'`）会被静默筛掉，判据于是把
    「登记 6 条」读成「登记 5 条」并反过来指控数据缩水。逐对解析后 key 与 reason 的对应
    关系才在判据里存在，才能对不同形态的 reason 用不同判据。
    """
    pairs: list[tuple[str, str]] = []
    pending: str | None = None
    for raw in body.splitlines():
        line = raw.strip()
        if not line:
            continue
        if pending is not None:
            value = re.match(r"'((?:[^'\\]|\\.)*)'", line)
            if value:
                pairs.append((pending, value.group(1)))
                pending = None
            continue
        head = re.match(r"(?:'((?:[^'\\]|\\.)*)'|([^:'\s][^:]*?))\s*:\s*(.*)$", line)
        if not head:
            continue
        key = head.group(1) if head.group(1) is not None else (head.group(2) or "").strip()
        rest = head.group(3).strip()
        if not rest:
            pending = key
            continue
        value = re.match(r"'((?:[^'\\]|\\.)*)'", rest)
        if value:
            pairs.append((key, value.group(1)))
    return pairs


def _strip_null_placeholder_columns(window: str) -> str:
    """把 `conclusion: null` / `remark: undefined` 这类**显式空占位**从写入点窗口剔除。

    🔴 `payload_column_mode` 判的是「业务内容写进哪一列」，而 `conclusion: null` 恰恰是
    「**不**往这列写内容」。G2 的写入点（`useG2Detail.ts#L517`）是
    `{ item_id, conclusion: null, remark: JSON.stringify(rows) }` —— 只看 token
    `conclusion:` 在不在窗口里，会把它误判成 `dual_write_remark_and_conclusion`。
    🔴 剔除只针对**字面 null / undefined**，真双写（`{ remark: json, conclusion: json }`，
    G4-9 / G5 / G6-main / G6-ecl 四条实测形态）一个都不会被剔掉；且哪条 entry 有空占位
    由 slice 的 `payload_null_placeholder_columns` 双向登记（声称有的源码里真有、没声称的
    真没有），见 `test_payload_null_placeholder_is_declared_both_ways`。
    """
    return re.sub(r"\b(remark|conclusion)\s*:\s*(?:null|undefined)\s*,?", "", window)


#: 位置化行身份的 token —— 命中即说明身份值由数组下标/序号派生（forbidden 三种之一）。
_POSITIONAL_IDENTITY_TOKEN = re.compile(
    r"\$\{\s*i\s*\}"  # `…-${i}`
    r"|\$\{\s*seq\s*\}"  # `…-${seq}`（seq 由调用点绑成 i+1，G6-sppi 实测形态）
    r"|\$\{\s*index\s*\}"
    r"|(?:\|\||\?\?)\s*i\s*\+\s*1"  # `?? i + 1` / `|| i + 1`
    r"|\bindexOf\("  # 用位置反查身份
)


def _value_expr_after_key(line: str, key: str) -> str | None:
    """截出 `key:` 之后**属于它自己**的值表达式（到同层逗号 / 行尾止）。

    🔴 为什么不能拿整行当判据：G12 的载入行是
    `enrich({ ...r, rowId: r.rowId || genId(), seq: r.seq ?? i + 1 })` —— 整行含
    `i + 1`，但那是 `seq`（展示序号，`resequence()` 每次重排都会重算）的值，不是 `rowId`
    的。按整行判会把 G12 误判成位置化行身份。反过来 G6-sppi 的位置化 token 就落在身份值
    自己的表达式内，截到同层逗号为止两者才分得开。
    """
    match = re.search(r"\b" + re.escape(key) + r"\s*\??\s*:", line)
    if not match:
        return None
    rest = line[match.end() :]
    depth = 0
    in_tmpl = False
    for pos, char in enumerate(rest):
        if char == "`":
            in_tmpl = not in_tmpl
            continue
        if in_tmpl:
            continue
        if char in "([{":
            depth += 1
        elif char in ")]}":
            if depth == 0:
                return rest[:pos]
            depth -= 1
        elif char == "," and depth == 0:
            return rest[:pos]
    return rest


def _split_top_level_args(arglist: str) -> list[str]:
    args: list[str] = []
    depth = 0
    in_tmpl = False
    buf = ""
    for char in arglist:
        if char == "`":
            in_tmpl = not in_tmpl
        if not in_tmpl:
            if char in "([{":
                depth += 1
            elif char in ")]}":
                depth -= 1
            elif char == "," and depth == 0:
                args.append(buf)
                buf = ""
                continue
        buf += char
    if buf.strip():
        args.append(buf)
    return args


def _functions_whose_first_param_is(source: str, key: str) -> set[str]:
    """同模块内首参名 == `key` 的函数名集合。"""
    names: set[str] = set()
    for match in re.finditer(r"function\s+([A-Za-z0-9_$]+)\s*\(\s*([A-Za-z0-9_$]+)", source):
        if match.group(2) == key:
            names.add(match.group(1))
    for match in re.finditer(
        r"const\s+([A-Za-z0-9_$]+)\s*=\s*(?:async\s*)?\(\s*([A-Za-z0-9_$]+)", source
    ):
        if match.group(2) == key:
            names.add(match.group(1))
    return names


def _positional_identity_hits(source: str, key: str) -> list[tuple[int, str]]:
    """行身份值是否由位置派生 —— 两类构造点各自取**身份自己的表达式**再判。

    (a) 对象属性形态 `<key>: <expr>`；
    (b) 同模块内**首参名就是身份键**的函数（如 `emptyRow(id, seq, name)`）在调用点的第 0
        个实参 —— 这是走一跳参数绑定，G6-sppi 的真缺陷正是这个形态
        （`emptyRow(String(raw.id || \\`fv-${Date.now()}-${seq}\\`), seq, …)`，身份键
        `id` 从来不以 `id:` 出现在对象字面量里）。而 G12 的 `migrateLegacyRow(raw, seq)`
        首参叫 `raw` 不叫 `rowId`，其调用点的 `i + 1` 绑给的是 `seq`，不进本判据 —— 这正是
        (b) 必须按**声明的参数名**而不是「附近有 migrate*Row(」来取的原因。
    """
    hits: list[tuple[int, str]] = []
    lines = source.splitlines()
    for no, line in enumerate(lines, 1):
        expr = _value_expr_after_key(line, key)
        if expr is not None and _POSITIONAL_IDENTITY_TOKEN.search(expr):
            hits.append((no, expr.strip()[:160]))
    for fname in _functions_whose_first_param_is(source, key):
        for match in re.finditer(r"\b" + re.escape(fname) + r"\s*\(", source):
            head = source[max(0, match.start() - 40) : match.start()]
            if re.search(r"(?:function|const|let|var)\s+$", head):
                continue  # 声明处，不是调用点
            tail = source[match.end() :]
            depth = 1
            in_tmpl = False
            end = None
            for pos, char in enumerate(tail):
                if char == "`":
                    in_tmpl = not in_tmpl
                    continue
                if in_tmpl:
                    continue
                if char in "([{":
                    depth += 1
                elif char in ")]}":
                    depth -= 1
                    if depth == 0:
                        end = pos
                        break
            if end is None:
                continue
            args = _split_top_level_args(tail[:end])
            if not args:
                continue
            if _POSITIONAL_IDENTITY_TOKEN.search(args[0]):
                no = source[: match.start()].count("\n") + 1
                hits.append((no, args[0].strip()[:160]))
    return hits


def _http_probe(client: str) -> Callable[[str], bool]:
    """按 slice 登记的 HTTP 客户端标识拼探针。

    🔴 不写成「`api.` 或 `http.` 二选一」的固定并集：那样任一 entry 的绑定名改了都不会红
    （另一族还在，并集恒真）。这里按**该 entry 自己声明的**标识拼正则 —— 声明与源码不一致
    立刻打红。
    """
    esc = re.escape(client)

    def probe(source: str) -> bool:
        return bool(
            re.search(rf"{esc}\.get\(\s*`[^`]*checklist-responses`", source)
        ) and bool(re.search(rf"{esc}\.put\(\s*`[^`]*checklist-responses`", source))

    return probe


# ────────────────────────────────────────────────────────────────────────────
# Fixtures
# ────────────────────────────────────────────────────────────────────────────
@pytest.fixture(scope="module")
def manifest_slice() -> dict:
    return _load(MANIFEST_SLICE_PATH)


@pytest.fixture(scope="module")
def deletion_plan() -> dict:
    return _load(DELETION_PLAN_PATH)


@pytest.fixture(scope="module")
def paradigm() -> dict:
    return _load(PARADIGM_PATH)


@pytest.fixture(scope="module")
def full_manifest() -> dict:
    return _load(FULL_MANIFEST_PATH)


# ════════════════════════════════════════════════════════════════════════════
# 判据自检 —— 判据本身不能恒真
# ════════════════════════════════════════════════════════════════════════════
class TestGuardSelfChecks:
    """反向自检：辅助函数既不能什么都不做，也不能把代码一起剥掉/截掉。"""

    def test_strip_comments_self_check(self) -> None:
        sample = (
            "// leading line comment\n"
            "/* block */\n"
            "const KEEP = 'value' // trailing\n"
            "<!-- html comment -->\n"
            "const ALSO = 'x'\n"
        )
        stripped = _strip_ts_comments(sample)
        for gone in ("leading line comment", "block", "trailing", "html comment"):
            assert gone not in stripped, f"{gone!r} 应被剥掉"
        assert "const KEEP = 'value'" in stripped
        assert "const ALSO = 'x'" in stripped

    def test_strip_comments_keeps_string_literals(self) -> None:
        """反向：剥注释不得吞掉字符串里的内容 —— 否则 G6 自造形态判据会恒真。"""
        sample = "const T = { label: '成本项目1' }\nconst U = 'generateRows'\n"
        stripped = _strip_ts_comments(sample)
        assert "成本项目1" in stripped, "字符串字面量被剥掉了 ⇒ 自造形态判据会恒真"
        assert "generateRows" in stripped

    def test_vue_template_extraction_does_not_stop_at_inner_template(self) -> None:
        """`</template>` 提前闭合的坑：截取必须以 `<script` 为界。"""
        sfc = (
            "<template>\n"
            "  <div>\n"
            "    <el-tooltip>\n"
            "      <template #content>inner</template>\n"
            "    </el-tooltip>\n"
            "    <MarkerTag />\n"
            "  </div>\n"
            "</template>\n"
            '<script setup lang="ts">\nconst x = 1\n</script>\n'
        )
        head = _vue_template(sfc)
        assert "<MarkerTag />" in head, (
            "外层模板尾部的标签被截掉了 —— 用 `</template>` 切会停在内层具名插槽处"
        )
        assert "const x = 1" not in head, "script 内容不该被当成模板"

    def test_toolbar_block_extraction_is_bounded(self) -> None:
        """区块截取必须按同缩进闭合，不能吞掉后面的兄弟节点。"""
        template = (
            "<template>\n"
            '  <div class="host">\n'
            '    <div class="g1-toolbar">\n'
            "      <A />\n"
            '      <div class="nested">\n'
            "        <B />\n"
            "      </div>\n"
            "    </div>\n"
            "    <OutsideTag />\n"
            "  </div>\n"
            "</template>\n"
        )
        block = _toolbar_block(template, 'class="g1-toolbar"')
        assert "<A />" in block and "<B />" in block, "区块内的节点被漏掉"
        assert "<OutsideTag />" not in block, "区块把兄弟节点吞进来了 ⇒ 判据范围失控"

    def test_toolbar_block_rejects_ambiguous_anchor(self) -> None:
        """锚点命中多处必须报错，而不是取第一处（取第一处会让判据锁错区块）。"""
        template = (
            "<template>\n"
            '  <div class="g1-toolbar">\n'
            "  </div>\n"
            '  <div class="g1-toolbar">\n'
            "  </div>\n"
            "</template>\n"
        )
        with pytest.raises(AssertionError, match="命中 2 次"):
            _toolbar_block(template, 'class="g1-toolbar"')

    def test_string_record_parser_keeps_short_values_and_multiline_values(self) -> None:
        """逐对解析器不得按长度筛掉短值，也要接住换行后的值。"""
        body = (
            "  甲表:\n"
            "    '源模板列为 A/B/C；底稿无这些字段',\n"
            "  '乙表（续：上年年末余额）': '同上（续表）',\n"
        )
        pairs = _parse_string_record(body)
        assert pairs == [
            ("甲表", "源模板列为 A/B/C；底稿无这些字段"),
            ("乙表（续：上年年末余额）", "同上（续表）"),
        ], f"逐对解析结果不对: {pairs}"

    def test_null_placeholder_stripper_only_removes_literal_nulls(self) -> None:
        """空占位剔除必须只吃字面 null/undefined，真双写一个都不能少。"""
        g2 = "{ item_id: KEY,\n  conclusion: null,\n  remark: JSON.stringify(rows),\n}"
        cleaned = _strip_null_placeholder_columns(g2)
        assert "remark:" in cleaned, "remark 被误剔 ⇒ remark_only 会判不出来"
        assert "conclusion:" not in cleaned, "conclusion: null 未被剔 ⇒ 会误判成双写"
        dual = "formData.debouncedSave(KEY, { remark: json, conclusion: json })"
        assert _strip_null_placeholder_columns(dual) == dual, (
            "真双写被剔掉了 ⇒ dual_write 形态会被误判成单写（假绿）"
        )

    def test_value_expr_after_key_is_scoped_to_its_own_key(self) -> None:
        """值表达式截取必须只取该键自己的部分（G12 与 G6-sppi 的分界线）。"""
        line = "      enrich({ ...r, rowId: r.rowId || genId(), seq: r.seq ?? i + 1 }),"
        assert _value_expr_after_key(line, "rowId") == " r.rowId || genId()", (
            "rowId 的值表达式被截到了 seq 那段 ⇒ G12 会被误判成位置化行身份"
        )
        assert "i + 1" in (_value_expr_after_key(line, "seq") or ""), (
            "seq 的值表达式截取失败 ⇒ 自检本身没有分母"
        )
        assert _value_expr_after_key(line, "rowKey") is None

    def test_positional_identity_detector_distinguishes_g12_from_g6_sppi(self) -> None:
        """两侧自检：参数绑定一跳要追到（G6-sppi 形态），序号列不得误伤（G12 形态）。"""
        g6_shape = (
            "function emptyRow(id: string, seq: number, name = ''): X {\n"
            "  return { id, seq, name }\n"
            "}\n"
            "export function migrateFairValueRow(raw: any, seq: number): X {\n"
            "  const base = emptyRow(String(raw.id || `fv-${Date.now()}-${seq}`), seq, '')\n"
            "  return base\n"
            "}\n"
            "const load = () => { rows.value = data.rows.map((r, i) => "
            "migrateFairValueRow(r, i + 1)) }\n"
        )
        hits = _positional_identity_hits(g6_shape, "id")
        assert hits, "参数绑定一跳的位置化身份没被追到 ⇒ G6-sppi 的真缺陷会漏报"
        assert any("${seq}" in snippet for _, snippet in hits), f"命中片段不对: {hits}"

        g12_shape = (
            "function migrateLegacyRow(raw: Record<string, unknown>, seq: number): Y {\n"
            "  return { rowId: String(raw.rowId ?? genId()), seq }\n"
            "}\n"
            "function enrich(raw: Partial<Y> & { rowId: string }): Y {\n"
            "  return { ...raw, rowId: raw.rowId, seq: parseNum(raw.seq) || 0 }\n"
            "}\n"
            "const load = (arr: any[]) => arr.map((r, i) =>\n"
            "  migrateLegacyRow(r, parseNum(r.seq) || i + 1),\n"
            ")\n"
            "const load2 = (arr: any[]) => arr.map((r, i) =>\n"
            "  enrich({ ...r, rowId: r.rowId || genId(), seq: r.seq ?? i + 1 }),\n"
            ")\n"
        )
        assert _positional_identity_hits(g12_shape, "rowId") == [], (
            "序号列 seq 的 i+1 被误当成行身份位置化 ⇒ G12 会假红"
        )
        # 反向：把身份真的接上下标就必须命中，否则本探针恒假
        g12_broken = g12_shape.replace(
            "rowId: String(raw.rowId ?? genId())", "rowId: `g12h-${seq}`"
        )
        assert _positional_identity_hits(g12_broken, "rowId"), (
            "把 rowId 直接接成 `${seq}` 后仍不命中 ⇒ 本探针恒假，判据无信息量"
        )

    def test_http_probe_is_client_specific(self) -> None:
        """探针必须按声明的客户端名拼，不能对另一族也放行。"""
        api_src = "const res = await api.get(`/api/workpapers/${x}/checklist-responses`)\n" \
                  "await api.put(`/api/workpapers/${x}/checklist-responses`, {})"
        http_src = "const { data } = await http.get(`/api/workpapers/${x}/checklist-responses`)\n" \
                   "await http.put(`/api/workpapers/${x}/checklist-responses`, {})"
        assert _http_probe("api")(api_src) is True
        assert _http_probe("api")(http_src) is False, "api 探针放行了 http 源码 ⇒ 绑定声明失去意义"
        assert _http_probe("http")(http_src) is True
        assert _http_probe("http")(api_src) is False
        only_get = "await api.get(`/api/workpapers/${x}/checklist-responses`)"
        assert _http_probe("api")(only_get) is False, "只读不写也应判否（业务内容不落 HTML 侧）"

    def test_all_required_artifacts_exist(self) -> None:
        """本文件依赖的每个真源都必须真存在 —— 缺真源一律 fail closed，不 skip。"""
        for path in (
            MANIFEST_SLICE_PATH,
            DELETION_PLAN_PATH,
            PARADIGM_PATH,
            FULL_MANIFEST_PATH,
            D_SLICE_PATH,
            E_SLICE_PATH,
            F_SLICE_PATH,
            CONTRACT_DIR,
            G7_PILOT_CONTRACT,
            TEMPLATE_DIR / "G",
            TEMPLATE_INDEX,
            CHECKLIST_ROUTER,
            TEMPLATE_FINDER,
            REGISTRY,
            NOTICE_MODULE,
            NOTICE_COMPONENT,
            HTML_RENDERER_REGISTRY,
            G1_SHEET_LABELS,
            G1_DUAL_MODE,
            E1_HOST,
            G6_NOTE_SECTION_MAP,
            G4_NOTE_SECTION_MAP,
            G6_LISTED_TAB,
            NOTE_SYNC_REGISTRY,
            UNREACHABLE_STUB,
            SHARED_BASE,
        ):
            assert path.exists(), (
                f"{path} 不存在。缺真源必须打红：`if not x.exists(): return` / "
                "`pytest.skip(...)` 都是 fail-open —— 文件一消失判据就自动通过"
            )


# ════════════════════════════════════════════════════════════════════════════
# AC 12.8 / 12.9 / 12.1 / 1.3：裁决合法性
# ════════════════════════════════════════════════════════════════════════════
class TestAdjudicationLegality:
    """
    **Validates: Requirements 12.1**

    AC 12.8 原文：「无 HTML 对端的纯 OO entry SHALL 标为 `single_onlyoffice`；
    不得为满足数字伪造字段映射或修改模板制造对端。」判据是**有没有 HTML 对端**，
    不是「adapter / contract / bundle 还不存在」。
    """

    def test_every_entry_has_a_binary_html_counterpart_verdict(
        self, manifest_slice: dict
    ) -> None:
        """step 3 必须产出二值结论；unresolved / unknown / 空值都不是结论（AP-3）。"""
        assert manifest_slice["independent_entries"], "entry 列表为空 ⇒ 后续判据全无分母"
        for entry in manifest_slice["independent_entries"]:
            verdict = entry.get("html_counterpart_verdict")
            assert verdict in HTML_COUNTERPART_VERDICTS, (
                f"{entry['entry_id']} 的 html_counterpart_verdict={verdict!r} 不是二值结论。"
                f"合法值只有 {HTML_COUNTERPART_VERDICTS}"
            )

    def test_single_onlyoffice_requires_no_html_counterpart(
        self, manifest_slice: dict
    ) -> None:
        """AC 12.8 的蕴含关系：single_onlyoffice ⇒ verdict == 'none'。"""
        for entry in manifest_slice["independent_entries"]:
            if entry.get("capability") == "single_onlyoffice":
                assert entry.get("html_counterpart_verdict") == "none", (
                    f"{entry['entry_id']} 裁为 single_onlyoffice 但 html_counterpart_verdict="
                    f"{entry.get('html_counterpart_verdict')!r}。AC 12.8 只允许「无 HTML 对端」"
                    "的 entry 标 single_onlyoffice"
                )

    def test_adjudication_reason_is_not_circular(
        self, manifest_slice: dict, paradigm: dict
    ) -> None:
        """AP-1：不得用「本任务应当交付的产物尚不存在」当裁决理由。

        判据取范式 JSON 登记的 `circular_reason_markers`（真源在范式，不在这里另写一份）。
        三个理由字段都扫 —— 只扫 `reason` 会让循环论证换到
        `not_bidirectional_because` 里照样通过。
        """
        markers = paradigm["adjudication_criteria"]["anti_patterns"][0][
            "circular_reason_markers"
        ]
        assert markers, "circular_reason_markers 为空 ⇒ 判据恒真"
        for entry in manifest_slice["independent_entries"]:
            adjudication = entry["adjudication"]
            for field in ("reason", "not_single_html_because", "not_bidirectional_because"):
                text = adjudication.get(field, "")
                hits = [m for m in markers if m in text]
                assert not hits, (
                    f"{entry['entry_id']} 的 adjudication.{field} 命中循环论证标记 {hits}。"
                    "缺什么写进 blocking_preconditions，不要当裁决依据"
                )

    def test_capability_is_enum_or_explicitly_pending(self, manifest_slice: dict) -> None:
        """AC 1.3：capability 只能是四个枚举值；当下没有一个成立时必须显式记 pending。

        🔴 这不是「允许留空」：capability 为 null 时必须同时给 `capability_verdict_stage`、
        `capability_target` 与非空 `capability_target_blocked_by`，并计入 unadjudicated。
        留空而不解释就是「未裁决伪装成已裁决」。
        """
        for entry in manifest_slice["independent_entries"]:
            capability = entry.get("capability")
            if capability is None:
                assert entry.get("capability_verdict_stage"), (
                    f"{entry['entry_id']} capability 为 null 却没有 capability_verdict_stage"
                )
                assert entry.get("capability_target") in CAPABILITY_ENUM, (
                    f"{entry['entry_id']} 的 capability_target 必须落在能力态枚举内"
                )
                blocked_by = entry.get("capability_target_blocked_by")
                assert isinstance(blocked_by, list) and blocked_by, (
                    f"{entry['entry_id']} capability 未定却没登记阻断项"
                )
            else:
                assert capability in CAPABILITY_ENUM, (
                    f"{entry['entry_id']} capability={capability!r} 不在 {CAPABILITY_ENUM}"
                )

    def test_capability_matches_honest_capability(self, manifest_slice: dict) -> None:
        """SR-4：对外 capability 与 adjudication.honest_capability 不得双口径。"""
        for entry in manifest_slice["independent_entries"]:
            assert entry.get("capability") == entry["adjudication"].get("honest_capability"), (
                f"{entry['entry_id']} capability 与 honest_capability 不一致"
            )

    def test_bidirectional_requires_all_five_identity_fields(
        self, manifest_slice: dict
    ) -> None:
        """SR-6 / AC 12.1：标 bidirectional 必须五个身份字段全非空。"""
        keys = (
            "adapter_id",
            "authority_model",
            "definition_bundle",
            "instrumentation_candidate",
            "published_representation",
        )
        for entry in manifest_slice["independent_entries"]:
            if entry.get("capability") == "bidirectional":
                for key in keys:
                    assert entry.get(key) is not None, (
                        f"{entry['entry_id']} 标 bidirectional 但 {key} 为空"
                    )

    def test_single_or_pending_entries_carry_no_identity(self, manifest_slice: dict) -> None:
        """SR-5 / AP-5：裁 single 或终态未定的 entry 不得挂身份字段（不伪造凑数）。"""
        keys = (
            "adapter_id",
            "authority_model",
            "definition_bundle",
            "instrumentation_candidate",
            "published_representation",
        )
        for entry in manifest_slice["independent_entries"]:
            capability = entry.get("capability")
            if capability == "bidirectional":
                continue
            for key in keys:
                assert entry.get(key) is None, (
                    f"{entry['entry_id']} capability={capability!r} 却挂了 {key}={entry[key]!r}"
                )

    def test_adjudication_carries_both_negative_reasons(self, manifest_slice: dict) -> None:
        """step 4 要求：not_single_html_because 与 not_bidirectional_because 都要写。"""
        for entry in manifest_slice["independent_entries"]:
            adjudication = entry["adjudication"]
            for key in ("not_single_html_because", "not_bidirectional_because"):
                assert len(adjudication.get(key, "")) > 40, (
                    f"{entry['entry_id']} 的 {key} 缺失或过短 —— 排除另外两个枚举值也要给理由"
                )

    def test_capability_blockers_reference_real_preconditions(
        self, manifest_slice: dict
    ) -> None:
        """entry 上登记的阻断项 id 必须真在 blocking_preconditions 里。"""
        known = {bp["id"] for bp in manifest_slice["blocking_preconditions"]}
        for entry in manifest_slice["independent_entries"]:
            for bp_id in entry.get("capability_target_blocked_by", []):
                assert bp_id in known, (
                    f"{entry['entry_id']} 引用了不存在的阻断项 {bp_id}"
                )

    def test_every_bidirectional_blocker_has_a_carrier(self, manifest_slice: dict) -> None:
        """反向：声称阻断 bidirectional 的阻断项必须被某 entry 引用（否则是死条目）。

        没有承载者的阻断项 = 死条目：它可以被修好而无人回来删，也可以永远挂着而不阻断
        任何东西。slice 级的（BP-6/BP-9/BP-11/BP-12）在 `blocks` 里明写不是 bidirectional。
        """
        referenced: set[str] = set()
        for entry in manifest_slice["independent_entries"]:
            referenced.update(entry.get("capability_target_blocked_by", []))
        for bp in manifest_slice["blocking_preconditions"]:
            if bp["id"] in referenced:
                continue
            assert bp["blocks"] != "bidirectional", (
                f"{bp['id']} 声称阻断 bidirectional 却没有任何 entry 引用它 ⇒ 死条目"
            )

    def test_manifest_mirror_divergence_is_registered_not_silently_equal(
        self, manifest_slice: dict, full_manifest: dict
    ) -> None:
        """BP-6：manifest 的 capability/html_store 是 overlay 组件级默认值，不是裁决结果。

        🔴 判据方向是「**必须**不一致且该不一致已登记」，不是「必须相等」。断言相等就是把
        overlay 默认值当成裁决真源。
        """
        by_id = {e["entry_id"]: e for e in full_manifest["entries"]}
        bp6 = next(bp for bp in manifest_slice["blocking_preconditions"] if bp["id"] == "BP-6")
        assert bp6["status"] == "REGISTERED_NOT_FIXED"
        for entry in manifest_slice["independent_entries"]:
            src = by_id[entry["entry_id"]]
            mirror = entry["manifest_mirror"]
            assert mirror["capability"] == src["capability"], (
                f"{entry['entry_id']} 的 manifest_mirror.capability 与 source manifest 不符"
            )
            assert mirror["html_store"] == src["html_store"], (
                f"{entry['entry_id']} 的 manifest_mirror.html_store 与 source manifest 不符"
            )
            assert src["capability"] != entry["capability"], (
                f"{entry['entry_id']}: source manifest 的 capability 与 slice 裁决相等了 ⇒ "
                "要么 overlay 已支持 per-entry override（请更新 BP-6），要么本 slice 把 overlay "
                "默认值当成了裁决结果"
            )
            assert src["html_store"] == "unresolved", (
                f"{entry['entry_id']}: manifest 的 html_store 不再是 overlay 默认值 "
                f"{src['html_store']!r} ⇒ BP-6 描述已过期"
            )
            assert entry["html_counterpart_verdict"] == "exists"


# ════════════════════════════════════════════════════════════════════════════
# HTML 对端事实的三边锁：slice ↔ 前端源码 ↔ 权威 xlsx
# ════════════════════════════════════════════════════════════════════════════
class TestHtmlCounterpartIsSourceBacked:
    """
    **Validates: Requirements 12.4**

    「有 HTML 对端」这个结论不能是自由文本 —— 每条都要能在磁盘上复核。
    G 循环的五处形态差异（payload 列四种 / HTTP 客户端两族 / 一册多 entry / 行身份三族 /
    键声明散落）逐 entry 从 slice 读声明再回源码核对，不写死单一形态。
    """

    def test_source_refs_point_at_real_paths(self, manifest_slice: dict) -> None:
        """每条 source_ref 的路径部分必须真存在（含 sheet!cell 形态的模板引用）。"""
        checked = 0
        for entry in manifest_slice["independent_entries"]:
            refs = entry.get("html_counterpart_source_refs")
            assert isinstance(refs, list) and len(refs) >= 7, (
                f"{entry['entry_id']} 的 html_counterpart_source_refs 太少（至少要有宿主 / "
                "持久化 composable / 主表 owner / 写入点 / 两个后端端点 / 模板单元格）"
            )
            for ref in refs:
                raw = ref.split("#")[0].split("!")[0].strip()
                candidate = ROOT / raw
                assert candidate.exists(), (
                    f"{entry['entry_id']} 的 source_ref {ref!r} 指向不存在的路径 {raw!r}"
                )
                checked += 1
        assert checked >= 119, f"source_ref 校验数 {checked} 太少 —— 分母可疑"

    def test_html_store_endpoints_exist_in_the_router(self, manifest_slice: dict) -> None:
        """slice 声称的读写端点必须真在 checklist_responses 路由里。"""
        router = CHECKLIST_ROUTER.read_text(encoding="utf-8")
        assert 'prefix="/api/workpapers/{wp_id}/checklist-responses"' in router
        assert '@router.get("", response_model=list[ChecklistResponseOut])' in router
        assert '@router.put("", response_model=list[ChecklistResponseOut])' in router
        assert "FROM checklist_responses" in router
        for entry in manifest_slice["independent_entries"]:
            store = entry["html_counterpart"]
            assert store["store"] == "checklist_responses"
            assert store["endpoint_read"].endswith("/checklist-responses")
            assert store["endpoint_write"].endswith("/checklist-responses")

    def test_payload_column_mode_is_declared_and_matches_the_write_site(
        self, manifest_slice: dict
    ) -> None:
        """FD-1：payload 列四种形态逐 entry 回**写入点源码**核验。

        🔴 F 循环的守卫写死 `payload_column == "remark"`。G 循环 17 条里只有 7 条如此；
        写死会让另 10 条判错。这里按 slice 声明的 mode 各用独立探针：
        * `remark_only`          ⇒ 写入点出现 `remark:` 且**不出现** `conclusion:` 同行双写；
        * `conclusion_only`      ⇒ 写入点出现 `conclusion:` 且不出现 `remark:`；
        * `dual_write_remark_and_conclusion` ⇒ 同一次调用里两列都出现；
        * `conclusion_canonical_remark_mirror` ⇒ 写入点走 storage contract 的
          `buildCanonicalPayload`，且该 contract 模块里 buildCanonicalPayload 真的同文双写。
        """
        allowed = {
            "remark_only": "remark",
            "conclusion_only": "conclusion",
            "dual_write_remark_and_conclusion": "conclusion",
            "conclusion_canonical_remark_mirror": "conclusion",
        }
        modes: dict[str, int] = {}
        for entry in manifest_slice["independent_entries"]:
            store = entry["html_counterpart"]
            mode = store.get("payload_column_mode")
            assert mode in allowed, (
                f"{entry['entry_id']} 的 payload_column_mode={mode!r} 不在 {sorted(allowed)}"
            )
            assert store["payload_column"] == allowed[mode], (
                f"{entry['entry_id']}: payload_column={store['payload_column']!r} 与 mode="
                f"{mode!r} 不一致（mode 决定权威列）"
            )
            modes[mode] = modes.get(mode, 0) + 1

            ref = store["payload_column_source"]
            path = _resolve_frontend(ref)
            assert path.is_file(), (
                f"{entry['entry_id']} 的 payload_column_source 不存在: {path}"
            )
            line_no = int(ref.rsplit("#L", 1)[1])
            raw_lines = path.read_text(encoding="utf-8").splitlines()
            assert 1 <= line_no <= len(raw_lines), (
                f"{entry['entry_id']} 的 payload_column_source 行号 {line_no} 越界"
            )
            # 写入点可能跨行（`debouncedSave(KEY, {` 换行），取声明行起 4 行窗口
            window = "\n".join(raw_lines[line_no - 1 : line_no + 3])
            window = _strip_ts_comments(window)
            # 🔴 先剔掉显式空占位再判：`conclusion: null` 是「不往这列写内容」，把它算成
            # 「写了 conclusion」会把 G2 的 remark_only 误判成 dual_write。
            probe_window = _strip_null_placeholder_columns(window)
            has_remark = "remark:" in probe_window
            has_conclusion = "conclusion:" in probe_window
            has_canonical = "buildCanonicalPayload(" in probe_window
            if mode == "remark_only":
                assert has_remark and not has_conclusion, (
                    f"{entry['entry_id']}: 声称 remark_only，但 {ref} 附近实测 "
                    f"remark={has_remark} conclusion={has_conclusion}\n{window}"
                )
            elif mode == "conclusion_only":
                assert has_conclusion and not has_remark, (
                    f"{entry['entry_id']}: 声称 conclusion_only，但 {ref} 附近实测 "
                    f"remark={has_remark} conclusion={has_conclusion}\n{window}"
                )
            elif mode == "dual_write_remark_and_conclusion":
                assert has_remark and has_conclusion, (
                    f"{entry['entry_id']}: 声称双写，但 {ref} 附近只见 "
                    f"remark={has_remark} conclusion={has_conclusion}\n{window}"
                )
            else:  # conclusion_canonical_remark_mirror
                assert has_canonical, (
                    f"{entry['entry_id']}: 声称走 storage contract 的 buildCanonicalPayload，"
                    f"但 {ref} 附近没有该调用\n{window}"
                )
                # 第二边：contract 模块里 buildCanonicalPayload 必须真的同文双写
                contract_module = None
                for candidate in COMPOSABLES.glob("g*StorageContract.ts"):
                    text = _strip_ts_comments(candidate.read_text(encoding="utf-8"))
                    if "export function buildCanonicalPayload" in text:
                        body = text.split("export function buildCanonicalPayload", 1)[1][:600]
                        if "conclusion: json" in body and "remark: json" in body:
                            contract_module = candidate
                            break
                assert contract_module is not None, (
                    f"{entry['entry_id']}: 找不到任何 g*StorageContract.ts 的 "
                    "buildCanonicalPayload 同文双写 conclusion + remark ⇒ mode 声明不成立"
                )
        assert len(modes) >= 4, (
            f"G 循环实测四种 payload 列形态并存，slice 只出现 {sorted(modes)} ⇒ 调查缩水"
        )
        declared = manifest_slice["honest_adjudication_summary"]["payload_column_mode_counts"]
        assert declared == modes, f"摘要的 payload 形态分布 {declared} 与逐 entry 现算 {modes} 不符"

        # 🔴 FD-1 的散文曾与机器计数三处不符（remark 单写写「6 条」却列了 7 个 wp_code 且漏了 G2；
        # storage contract 那族把 G4-sppi 算进去，而它的写入点其实是逐字双写；双写族写「5 条」
        # 却列了 4 个）。散文不受判据管 ⇒ 数字必须搬进可复核字段。这里逐 entry 现算分组，与
        # FD-1 的 entries_by_mode 逐组比对。
        grouped: dict[str, list[str]] = {}
        for entry in manifest_slice["independent_entries"]:
            grouped.setdefault(
                entry["html_counterpart"]["payload_column_mode"], []
            ).append(entry["entry_id"])
        fd1 = next(
            d for d in manifest_slice["g_cycle_form_differences"]["differences"] if d["id"] == "FD-1"
        )
        declared_groups = fd1["entries_by_mode"]
        assert {k: sorted(v) for k, v in declared_groups.items()} == {
            k: sorted(v) for k, v in grouped.items()
        }, f"FD-1 的 entries_by_mode 与逐 entry 现算分组不符：\n登记 {declared_groups}\n现算 {grouped}"

    def test_payload_null_placeholder_is_declared_both_ways(
        self, manifest_slice: dict
    ) -> None:
        """FD-1 的第二半：显式空占位列必须逐 entry 登记，且双向可复核。

        🔴 `_strip_null_placeholder_columns` 让上一条判据能区分「写 null」和「写内容」，
        但它同时也是一把**放宽**判据的刀 —— 如果没人管住「谁有空占位」，将来某条 entry 真的
        退化成 `conclusion: <内容>` 也不会有人发现（剔除规则只吃字面 null，但登记若缺失，
        「窗口里出现 conclusion: null」这件事就从此无人核对）。所以两个方向都锁：
        * 声称有空占位（`payload_null_placeholder_columns` 非空）⇒ 写入点窗口里真有
          `<列>: null`；
        * 没声称（空数组）⇒ 窗口里真的一个字面 null 列都没有。
        全 slice 实测只有 G2 一条有空占位（`conclusion: null`）。
        """
        with_placeholder: dict[str, list[str]] = {}
        for entry in manifest_slice["independent_entries"]:
            store = entry["html_counterpart"]
            declared = store.get("payload_null_placeholder_columns")
            assert isinstance(declared, list), (
                f"{entry['entry_id']} 缺 payload_null_placeholder_columns（空占位列表，"
                "没有就写 []，不允许缺键 —— 缺键会让上面那条剔除规则无人核对）"
            )
            assert set(declared) <= {"remark", "conclusion"}, (
                f"{entry['entry_id']} 的 payload_null_placeholder_columns={declared} "
                "含非法列名（只可能是 remark / conclusion）"
            )
            ref = store["payload_column_source"]
            path = _resolve_frontend(ref)
            line_no = int(ref.rsplit("#L", 1)[1])
            raw_lines = path.read_text(encoding="utf-8").splitlines()
            window = _strip_ts_comments("\n".join(raw_lines[line_no - 1 : line_no + 3]))
            actual = sorted(
                set(re.findall(r"\b(remark|conclusion)\s*:\s*(?:null|undefined)\b", window))
            )
            assert actual == sorted(declared), (
                f"{entry['entry_id']}: payload_null_placeholder_columns 声称 "
                f"{sorted(declared)}，{ref} 附近实测 {actual}\n{window}"
            )
            if declared:
                with_placeholder[entry["entry_id"]] = sorted(declared)
        assert with_placeholder == {"xlsx/gt-g2-interest-receivable": ["conclusion"]}, (
            f"带显式空占位的 entry 实测为 {with_placeholder}，与冻结结论不符"
        )
        summary = manifest_slice["honest_adjudication_summary"]
        assert summary["payload_null_placeholder_entries"] == len(with_placeholder), (
            f"摘要的空占位 entry 数 {summary['payload_null_placeholder_entries']} "
            f"与现算 {len(with_placeholder)} 不符"
        )

    def test_persistence_composable_really_calls_both_endpoints_with_its_own_client(
        self, manifest_slice: dict
    ) -> None:
        """第二边（持久化侧）：声称的 composable 必须用**它自己声明的**客户端 GET 且 PUT。

        🔴 FD-2：G 循环有两族 HTTP 客户端。F 的探针写死 `api.get(...)`，对 4 条用 `http` 的
        entry 必然失配。这里按 slice 登记的 `http_client_binding` 拼探针，并额外核验 import
        语句逐字存在 —— 声明与源码不一致立刻打红。
        """
        clients: dict[str, int] = {}
        for entry in manifest_slice["independent_entries"]:
            store = entry["html_counterpart"]
            declared = store["html_persistence_composable"]
            module = _resolve_frontend(declared)
            assert module.is_file(), (
                f"{entry['entry_id']} 的 html_persistence_composable 不存在: {module}"
            )
            source = _strip_ts_comments(module.read_text(encoding="utf-8"))
            client = store["http_client_binding"]
            clients[client] = clients.get(client, 0) + 1
            import_line = store["http_client_import"]
            assert import_line in source, (
                f"{entry['entry_id']}: {module.name} 里找不到声明的 import 语句 {import_line!r}"
            )
            assert _http_probe(client)(source), (
                f"{entry['entry_id']}: {module.name} 里找不到 {client}.get + {client}.put "
                "两个 checklist-responses 调用 ⇒ 只读不写的话业务内容并不落在 HTML 侧，"
                "或客户端绑定声明写错了"
            )
        assert len(clients) == 2, (
            f"G 循环实测两族 HTTP 客户端并存，slice 只出现 {sorted(clients)} ⇒ 调查缩水"
        )
        declared_counts = manifest_slice["honest_adjudication_summary"][
            "http_client_binding_counts"
        ]
        assert declared_counts == clients, (
            f"摘要的客户端分布 {declared_counts} 与逐 entry 现算 {clients} 不符"
        )

    def test_transport_key_kind_is_declared_and_matches_the_source(
        self, manifest_slice: dict
    ) -> None:
        """传输键声明形态逐 entry 回源码核验 —— 四种形态各有独立判据。

        * `literal_item_id`                        ⇒ owner 模块里有 `const X = '<item>'` 且常量名一致；
        * `literal_item_id_duplicated_across_modules` ⇒ 同上，且**另有**至少一个模块也写了同值
          字面量（这条形态本身就是 BP-10 的实例，必须能复现）；
        * `storage_contract_member_item_id`        ⇒ owner 是 g*StorageContract.ts，键以
          `G{n}_ITEM_IDS` 成员形式声明；
        * `per_sheet_literal_declared_in_the_consuming_sfc` ⇒ owner 是 `.vue`，键在 SFC 内
          `const X = '<item>'`。
        """
        allowed = {
            "literal_item_id",
            "literal_item_id_duplicated_across_modules",
            "storage_contract_member_item_id",
            "per_sheet_literal_declared_in_the_consuming_sfc",
        }
        kinds: dict[str, int] = {}
        for entry in manifest_slice["independent_entries"]:
            store = entry["html_counterpart"]
            kind = store.get("transport_key_kind")
            assert kind in allowed, (
                f"{entry['entry_id']} 的 transport_key_kind={kind!r} 不在 {sorted(allowed)}"
            )
            kinds[kind] = kinds.get(kind, 0) + 1
            item_id = store["transport_key_shape"]
            assert item_id == store["primary_table"]["item_id"], (
                f"{entry['entry_id']}: transport_key_shape 与 primary_table.item_id 不一致"
            )
            owner = ROOT / store["primary_table"]["owner_module"]
            assert owner.is_file(), f"{entry['entry_id']} 的 owner_module 不存在: {owner}"
            source = _strip_ts_comments(owner.read_text(encoding="utf-8"))
            const_name = store["primary_table"]["owner_constant"]

            if kind in ("literal_item_id", "literal_item_id_duplicated_across_modules"):
                found = re.search(
                    r"const\s+([A-Z0-9_]+)\s*=\s*['\"]" + re.escape(item_id) + r"['\"]", source
                )
                assert found, (
                    f"{entry['entry_id']}: {owner.name} 里找不到 "
                    f"const {const_name} = '{item_id}'"
                )
                assert found.group(1) == const_name, (
                    f"{entry['entry_id']}: 常量名应为 {const_name}，实为 {found.group(1)}"
                )
            elif kind == "storage_contract_member_item_id":
                assert owner.name.endswith("StorageContract.ts"), (
                    f"{entry['entry_id']}: 声称 storage contract 成员键，但 owner 是 {owner.name}"
                )
                member = const_name.split(".", 1)[1]
                assert re.search(
                    rf"{re.escape(member)}\s*:\s*['\"]" + re.escape(item_id) + r"['\"]", source
                ), f"{entry['entry_id']}: {owner.name} 里找不到 {member}: '{item_id}'"
                assert const_name.split(".", 1)[0] in source, (
                    f"{entry['entry_id']}: {owner.name} 里找不到 "
                    f"{const_name.split('.', 1)[0]} 对象声明"
                )
            else:  # per_sheet_literal_declared_in_the_consuming_sfc
                assert owner.suffix == ".vue", (
                    f"{entry['entry_id']}: 声称 SFC 内声明，但 owner 是 {owner.name}"
                )
                assert re.search(
                    r"const\s+" + re.escape(const_name) + r"\s*=\s*['\"]"
                    + re.escape(item_id) + r"['\"]",
                    source,
                ), f"{entry['entry_id']}: {owner.name} 里找不到 const {const_name} = '{item_id}'"
        assert len(kinds) >= 3, (
            f"G 循环实测至少三种传输键声明形态并存，slice 只出现 {sorted(kinds)} ⇒ 调查缩水"
        )
        declared = manifest_slice["honest_adjudication_summary"]["transport_key_kind_counts"]
        assert declared == kinds, f"摘要的传输键形态分布 {declared} 与逐 entry 现算 {kinds} 不符"

    def test_duplicated_item_id_literal_scale_is_recomputable(
        self, manifest_slice: dict
    ) -> None:
        """BP-10：「80 个 item_id 被多模块各写一份」必须能现扫复算，不是自由文本。

        🔴 这条既锁规模（数字不是猜的），也锁修复后必须回来改 slice：把键收敛到 storage
        contract 后计数会下降 ⇒ 打红，逼作者更新 BP-10 的 status。
        """
        decl = re.compile(r"(?:const\s+[A-Za-z0-9_]+\s*=|[A-Z0-9_]+\s*:)\s*['\"](G\d+[A-Za-z0-9\-]*)['\"]")
        owners: dict[str, set[str]] = {}
        for path in FRONTEND.rglob("*.ts"):
            rel = path.relative_to(ROOT).as_posix()
            if "__tests__" in rel or "G7" in path.name or "/g7" in rel.lower():
                continue
            text = _strip_ts_comments(path.read_text(encoding="utf-8", errors="replace"))
            for match in decl.finditer(text):
                owners.setdefault(match.group(1), set()).add(rel)
        duplicated = {k: v for k, v in owners.items() if len(v) > 1}
        assert owners, "一个 G 前缀 item_id 声明都没扫到 ⇒ 正则失配，判据恒真"
        bp10 = next(bp for bp in manifest_slice["blocking_preconditions"] if bp["id"] == "BP-10")
        assert bp10["status"] == "REGISTERED_NOT_FIXED"
        declared = manifest_slice["honest_adjudication_summary"][
            "duplicated_item_id_literals_in_g_cycle"
        ]
        assert len(duplicated) == declared, (
            f"多模块重复声明的 item_id 实测 {len(duplicated)} 个，slice 登记 {declared} 个。"
            "若已开始收敛，请更新 BP-10 的 status 与本计数"
        )
        # 正向：G6 的派生别名不算重复字面量（它引用真源而非重写）
        cross = _strip_ts_comments((COMPOSABLES / "g6CrossHelpers.ts").read_text(encoding="utf-8"))
        assert "G6_11_ROWS_KEY = G6_ITEM_IDS.G6_11_ROWS" in cross, (
            "g6CrossHelpers 的派生别名形态已变 —— BP-10 里「G6 做对了」这个正向对照失效"
        )
        assert "G6-11-rows" not in cross, (
            "g6CrossHelpers 里出现了 'G6-11-rows' 字面量 ⇒ 派生别名退化成重复字面量"
        )

    def test_row_identity_key_and_generator_are_source_backed(
        self, manifest_slice: dict
    ) -> None:
        """行身份键与生成器逐 entry 回源码核对。

        🔴 FD-4：不能像 F 循环那样写死 `row_identity_key in ("rowId", "id")` —— G14 用的是
        源模板固定的 `rowKey`（行集由 G14_LINE_ITEMS 决定，用户不增删），那是**最稳**的一族，
        写死两族会把它判非法。这里三族各有独立判据。
        """
        keys: set[str] = set()
        kinds: dict[str, int] = {}
        for entry in manifest_slice["independent_entries"]:
            store = entry["html_counterpart"]
            key = store["row_identity_key"]
            assert key in ("rowId", "id", "rowKey"), (
                f"{entry['entry_id']} 的 row_identity_key={key!r} 不是实测到的三族之一"
            )
            keys.add(key)
            kind = store["row_identity_kind"]
            kinds[kind] = kinds.get(kind, 0) + 1
            forbidden = store["forbidden_row_identity_kinds"]
            assert {"array_index", "ordinal", "position"} <= set(forbidden), (
                f"{entry['entry_id']} 的 forbidden_row_identity_kinds 缺位置化种类"
            )

            gen_ref = store["row_identity_generator_source"]
            gen_path = _resolve_frontend(gen_ref)
            assert gen_path.is_file(), (
                f"{entry['entry_id']} 的 row_identity_generator_source 不存在: {gen_path}"
            )
            gen_source = _strip_ts_comments(gen_path.read_text(encoding="utf-8"))
            # 身份字段必须真是该模块行模型里声明的字段
            assert re.search(rf"(?m)^\s*{key}\??:\s*string", gen_source), (
                f"{entry['entry_id']}: {gen_path.name} 的行模型里没有 `{key}: string` 声明 ⇒ "
                f"slice 声称行身份键是 {key!r}，但源码里没这个字段"
            )
            form = store["row_identity_generator_form"]
            if kind == "stable_template_row_key":
                assert "createDefaultRows" in gen_source, (
                    f"{entry['entry_id']}: 声称源模板固定行集，但 {gen_path.name} 里没有 "
                    "createDefaultRows"
                )
                assert "LINE_ITEMS" in gen_source, (
                    f"{entry['entry_id']}: 声称源模板固定行集，但找不到 *_LINE_ITEMS 行集常量"
                )
                assert "Date.now" not in form and "randomUUID" not in form, (
                    f"{entry['entry_id']}: 固定行集不该有生成器形态 {form!r}"
                )
                continue
            # 🔴 生成型三族统一判据：**声明的生成器形态必须逐字出现在源码里**（空白折叠后
            # 比对，因为源码会换行）。
            # 为什么不再按 kind 写死拼写探针：`generated_uuid` 曾写死 `"crypto.randomUUID"
            # in gen_source`，而 G6-ecl 的真实拼写是 `globalThis.crypto?.randomUUID?.()`
            # —— 可选链一插，写死的子串就失配，判据把「拼写不同」误报成「没有 uuid 生成」。
            # 换成逐字回源后两侧都咬：源码改了形态即红，slice 转写不精确（当初 G6-ecl 的
            # form 漏了 `globalThis.` 与 `crypto?.` 的问号）也即红。
            collapsed_form = re.sub(r"\s+", " ", form).strip()
            collapsed_source = re.sub(r"\s+", " ", gen_source)
            assert collapsed_form in collapsed_source, (
                f"{entry['entry_id']}: slice 声明的生成器形态 {form!r} 在 {gen_path.name} 里"
                "逐字找不到（已折叠空白）⇒ 声明与源码不符"
            )
            if kind == "generated_uuid":
                assert re.search(r"crypto\s*\??\s*\.\s*randomUUID", gen_source), (
                    f"{entry['entry_id']}: 声称 uuid 生成，但 {gen_path.name} 里没有 "
                    "crypto.randomUUID（含可选链拼写）"
                )
            else:
                probe = re.search(r"`([A-Za-z0-9]+)-\$\{", form)
                assert probe, (
                    f"{entry['entry_id']}: 生成器形态 {form!r} 里提不出前缀 ⇒ 无法回源码核对"
                )
                marker = f"`{probe.group(1)}-${{"
                assert marker in gen_source, (
                    f"{entry['entry_id']}: {gen_path.name} 里找不到生成器前缀 {marker!r}"
                )
        assert keys == {"rowId", "id", "rowKey"}, (
            f"G 循环实测三族行身份键并存，slice 只出现 {sorted(keys)} ⇒ 调查缩水或形态漂移"
        )
        declared = manifest_slice["honest_adjudication_summary"]["row_identity_kind_counts"]
        assert declared == kinds, f"摘要的行身份形态分布 {declared} 与逐 entry 现算 {kinds} 不符"

    def test_positional_row_identity_defects_are_real_and_exhaustive(
        self, manifest_slice: dict
    ) -> None:
        """位置化行身份缺陷必须双向锁死：声称有的源码里真有，声称没有的真没有。

        🔴 只查「声称有的真有」会让漏报（某 entry 也有下标派生但 slice 写 false）通过；
        只查「声称没有的真没有」会让虚报（拿一个不存在的缺陷凑 BP）通过。两侧都要。
        """
        flagged: list[str] = []
        for entry in manifest_slice["independent_entries"]:
            store = entry["html_counterpart"]
            gen_path = _resolve_frontend(store["row_identity_generator_source"])
            source = _strip_ts_comments(gen_path.read_text(encoding="utf-8"))
            id_key = store["row_identity_key"]
            # 🔴 判据从「行级」升到「表达式级 + 参数绑定一跳」，理由见
            # `_positional_identity_hits` 的 docstring：按行判会把 G12 的
            # `enrich({ ...r, rowId: r.rowId || genId(), seq: r.seq ?? i + 1 })` 判成位置化
            # （`i + 1` 是 seq 的值，不是 rowId 的），同时按「附近有 migrate*Row(」判又只是
            # 碰巧命中 G6-sppi —— 两者都不是「身份值由位置派生」这件事本身。
            hits = _positional_identity_hits(source, id_key)
            actual = bool(hits)
            declared = bool(store.get("row_identity_is_positional"))
            assert declared == actual, (
                f"{entry['entry_id']}: row_identity_is_positional 声称 {declared}，"
                f"但 {gen_path.name} 的行身份构造表达式实测 {actual}"
                f"（命中：{[f'#L{no} {snippet}' for no, snippet in hits]}）"
            )
            if declared:
                assert store.get("row_identity_positional_defect"), (
                    f"{entry['entry_id']} 声称位置化却没写 row_identity_positional_defect"
                )
                flagged.append(entry["entry_id"])
        assert set(flagged) == {"xlsx/gt-g6-other-bond-sppi"}, (
            f"位置化缺陷 entry 集合实测为 {sorted(flagged)}，与冻结结论不符"
        )
        # 另一侧：该缺陷的载入路径逐字可复现
        sppi = _strip_ts_comments(
            (COMPOSABLES / "useG6SppiFairValue.ts").read_text(encoding="utf-8")
        )
        assert "data.rows.map((r, i) => migrateFairValueRow(r, i + 1))" in sppi, (
            "BP-7 的载入路径已变 —— 若已修好请更新 slice 与本判据"
        )
        assert "`fv-${Date.now()}-${seq}`" in sppi, (
            "BP-7 的下标派生表达式已变 —— 若已修好请更新 slice 与本判据"
        )

    def test_primary_table_identity_cell_matches_the_authoritative_template(
        self, manifest_slice: dict
    ) -> None:
        """第三边：primary table 的身份列表头必须与权威模板单元格逐字相等。"""
        checked = 0
        for entry in manifest_slice["independent_entries"]:
            primary = entry["html_counterpart"]["primary_table"]
            template = TEMPLATE_DIR / primary["template_relative_path"]
            assert template.is_file(), f"{entry['entry_id']} 的权威模板不存在: {template}"
            actual = _cell_value(template, primary["template_sheet"], primary["identity_cell"])
            expected = primary["identity_text"]
            assert actual is not None and str(actual).strip() == expected, (
                f"{entry['entry_id']}: {primary['template_sheet']}!"
                f"{primary['identity_cell']} 实测 {actual!r}，slice 记的是 {expected!r}"
            )
            checked += 1
        assert checked == 17, f"只核验了 {checked} 组 sheet!cell，应为 17"

    def test_template_ref_resolves_through_the_runtime_index(
        self, manifest_slice: dict
    ) -> None:
        """template_ref 必须是运行时 finder 能拿到的文件（不在 _index.json 里的不算权威）。

        🔴 Task 46 的 D4 `D4收入底稿.xlsx` 与 Task 48 的 F2 `F2存货.xlsx` 都是「磁盘有、索引无、
        finder 永不返回」的同型陷阱。G 循环实测**没有**这种对象（15/15 全在索引里），本判据
        因此是双向锁：将来谁往 G 目录扔一本未索引的册子并写进 template_ref 就会打红。
        """
        indexed = _indexed_relpaths()
        assert len(indexed) > 400, f"_index.json 只索引到 {len(indexed)} 个文件，可疑"
        for entry in manifest_slice["independent_entries"]:
            ref = entry.get("template_ref")
            assert isinstance(ref, str) and ref.strip(), (
                f"{entry['entry_id']} 的 template_ref 为空 —— 不允许缺省跳过"
            )
            assert ref in indexed, (
                f"{entry['entry_id']} 的 template_ref {ref!r} 不在 "
                "backend/wp_templates/_index.json 里 ⇒ 运行时 finder 永不返回它，"
                "不能当权威模板"
            )
            assert (TEMPLATE_DIR / ref).is_file(), (
                f"{entry['entry_id']} 的 template_ref {ref!r} 在索引里但磁盘上不存在"
            )


# ════════════════════════════════════════════════════════════════════════════
# Property 20 / 21：contract 与 adapter —— **不宣称通过**的部分写在类名与 docstring 里
# ════════════════════════════════════════════════════════════════════════════
class TestProperty20And21NotClaimedPassingForThisSlice:
    """
    **Validates: Requirements 6.1, 6.2**

    ═══ 空分母怎么处理 ═══

    「本 slice 里没有 contract，所以 Property 20 / 21 通过」是重言式（同 spec 的 Task 20
    专门写代码拒绝这种论证）。诚实的写法是区分三件事：

    * **本 slice 无适用分母的部分** —— 本 slice 17 条 entry 已发布的 per-entry contract 数
      = 0（可复核事实：逐文件读 `review.entry_id`），**不宣称 Property 20 / 21 在这 17 条
      entry 上通过**，只断言前提成立且承载者真存在。
    * **本循环内有真分母的部分** —— G 循环有一份真实生产契约
      `g7.soe_subsidiary_disclosure.json`（G7 pilot 产物，entry_id 属**已排除**的 G7）。
      本类在它上面做 Property 20 的**正向**核验：逐字段有非空 stable key + source_ref，
      且没有 `col_[a-z]+` 无语义占位。这不是把 G7 算进本 slice，而是「本循环唯一能承载
      字段级判据的实体」——把它写清楚，后来者才知道 G 循环的 contract 长什么样。
    * **有真分母的其余部分** —— 见 Property 28 / 69 / 70 三个类。
    """

    def test_no_slice_entry_has_a_contract_and_the_g7_pilot_is_the_only_g_contract(
        self, manifest_slice: dict
    ) -> None:
        """分母为 0 这件事本身要被证实，并证明 G7 那条契约确实**不属**本 slice。"""
        slice_ids = {e["entry_id"] for e in manifest_slice["independent_entries"]}
        assert slice_ids, "slice entry 集合为空 ⇒ 本判据无分母"

        contract_files = sorted(CONTRACT_DIR.glob("*.json"))
        assert len(contract_files) >= 4, (
            f"契约目录只有 {len(contract_files)} 个文件 —— 分母可疑，"
            "「本 slice 没有契约」这个结论必须建立在真的读过全部契约之上"
        )
        owners: dict[str, str] = {}
        for path in contract_files:
            if path.name.startswith("_"):
                continue  # _example.candidate.json 是模板样例，不是生产契约
            entry_id = (_load(path).get("review") or {}).get("entry_id")
            if entry_id:
                owners[path.name] = entry_id
        assert owners, "读不出任何契约的 review.entry_id ⇒ 判据失去分母"

        leaked = {name: eid for name, eid in owners.items() if eid in slice_ids}
        assert not leaked, (
            f"契约 {leaked} 归属本 slice 的 entry，「本 slice contract 数 = 0」这个前提不再成立 "
            "⇒ 必须在此补齐字段级判据（stable_field_key / json_pointer / mode / "
            "value_type / source_ref 与 col_ 占位拒绝）"
        )
        # G 循环内确实有一条契约，且它属**已排除**的 G7 —— 这条前提必须为真，
        # 否则「本循环唯一能承载字段级判据的实体」这句话没有分母。
        g_owners = {n: e for n, e in owners.items() if re.match(r"^xlsx/(gt-)?g\d", e)}
        assert g_owners, (
            "契约目录里一条 G 循环契约都没有 ⇒ 下面的 Property 20 正向核验将无分母，"
            "本类会退化成纯重言式"
        )
        assert set(g_owners.values()) <= G7_ENTRY_IDS, (
            f"G 循环契约的归属 {g_owners} 不全在 G7 三条排除项里 ⇒ 本 slice 的排除边界要重算"
        )
        excluded_text = json.dumps(
            manifest_slice["slice_scope"]["excluded_from_slice"], ensure_ascii=False
        )
        for eid in g_owners.values():
            assert eid in excluded_text, (
                f"契约归属 {eid} 未在 excluded_from_slice 里显式排除 ⇒ 排除边界不可复核"
            )

    def test_g7_pilot_contract_has_no_generated_col_placeholder(self) -> None:
        """Property 20 的正向核验（真分母）：生产契约里不得有 `col_[a-z]+` 无语义占位。

        🔴 这不是「本 slice 通过」，而是「本循环唯一那份生产契约在这条 Property 上成立」。
        它同时锁死一件事：将来给本 slice 的 entry 发契约时，形态必须与它一致。
        """
        contract = _load(G7_PILOT_CONTRACT)
        # 🔴 `review_status` 在契约的**顶层**（与 contract_id / sheets 同级），不在 `review`
        # 子对象里 —— 全部 4 份生产契约（b60 / d2 / g7 / h1）实测都是这个形态。本条曾读
        # `review["review_status"]` 恒得 None 并当成「未审核」打红：那是判据读错路径，
        # 不是契约缺审核。两处都核对，防将来 schema 挪位后一侧静默变恒真。
        status = contract.get("review_status")
        nested = (contract.get("review") or {}).get("review_status")
        assert status == "reviewed", (
            f"g7 pilot 契约顶层 review_status={status!r}，"
            "非 reviewed 的契约不得注册生产 adapter（step 6 约束）"
        )
        assert nested in (None, "reviewed"), (
            f"契约 review.review_status={nested!r} 与顶层 {status!r} 不一致 ⇒ 审核态有两个真源"
        )
        placeholder = re.compile(r"^col_[a-z]+$")
        fields = 0
        for sheet in contract["sheets"]:
            for table in sheet["tables"]:
                for field in table["fields"]:
                    fields += 1
                    key = field.get("stable_field_key")
                    assert isinstance(key, str) and key.strip(), (
                        f"g7 契约里有字段缺 stable_field_key: {field}"
                    )
                    assert not placeholder.match(key), (
                        f"g7 契约出现无语义列占位 stable_field_key={key!r} ⇒ Property 20 违规"
                    )
                    src = field.get("source_ref")
                    assert isinstance(src, str) and src.strip(), (
                        f"g7 契约字段 {key!r} 缺 source_ref ⇒ Property 21 违规"
                    )
        assert fields >= 5, f"g7 契约只解析出 {fields} 个字段 ⇒ 分母可疑，判据会近似恒真"

    def test_property_21_carriers_exist(self) -> None:
        """「由 pilot 承载」不能是空头承诺 —— 承载文件必须真存在。"""
        for guard in (
            BACKEND / "tests" / "workpaper_sync" / "test_task13_contract_registry.py",
            BACKEND / "tests" / "workpaper_sync" / "test_task41_d2_large_json_pilot.py",
        ):
            assert guard.is_file(), (
                f"{guard} 不存在 —— 承载 Property 20/21 字段级判据的守卫缺失，"
                "「由 pilot 承载」就成了空头承诺"
            )

    def test_no_slice_entry_has_a_registered_adapter(
        self, manifest_slice: dict, full_manifest: dict
    ) -> None:
        """本 slice adapter 数 = 0 —— slice 与 source manifest 两侧都验。"""
        slice_ids = {e["entry_id"] for e in manifest_slice["independent_entries"]}
        for entry in manifest_slice["independent_entries"]:
            assert entry["adapter_id"] is None, (
                f"{entry['entry_id']} 出现了 adapter_id={entry['adapter_id']!r}，"
                "本 slice 的「无已注册 adapter」前提不再成立"
            )
        seen = 0
        for entry in full_manifest["entries"]:
            if entry["entry_id"] in slice_ids:
                assert entry.get("adapter_id") is None, (
                    f"source manifest 里 {entry['entry_id']} 已有 adapter_id"
                )
                seen += 1
        assert seen == len(slice_ids), (
            f"source manifest 里只找到 {seen}/{len(slice_ids)} 条 slice entry ⇒ slice 与 manifest 脱钩"
        )

    def test_registry_delivered_contracts_contain_no_slice_entry(
        self, manifest_slice: dict
    ) -> None:
        """registry 的 DELIVERED_PER_ENTRY_CONTRACTS 登记表里不得出现本 slice 的 entry。"""
        source = REGISTRY.read_text(encoding="utf-8")
        block = re.search(
            r"DELIVERED_PER_ENTRY_CONTRACTS[^=]*=\s*\(([\s\S]*?)\n\)\n", source
        )
        assert block, "找不到 DELIVERED_PER_ENTRY_CONTRACTS 的声明 ⇒ 判据无分母"
        declared = re.findall(r"\"entry_id\":\s*\"([^\"]+)\"", block.group(1))
        assert len(declared) >= 4, (
            f"登记表只解析出 {len(declared)} 条 entry_id ⇒ 正则失配，判据会恒真"
        )
        slice_ids = {e["entry_id"] for e in manifest_slice["independent_entries"]}
        leaked = [e for e in declared if e in slice_ids]
        assert not leaked, f"DELIVERED_PER_ENTRY_CONTRACTS 里出现本 slice 的 entry: {leaked}"
        # 反向：G7 那条**应当**在登记表里（它是 pilot）—— 若不在，说明我们读错了表
        g_declared = [e for e in declared if re.match(r"^xlsx/(gt-)?g\d", e)]
        assert set(g_declared) <= G7_ENTRY_IDS, (
            f"登记表里的 G 循环条目 {g_declared} 不全属 G7 ⇒ 排除边界要重算"
        )

    def test_property_denominator_block_declares_what_is_not_claimed(
        self, manifest_slice: dict
    ) -> None:
        """slice 必须显式写明哪几条 Property「不宣称通过」。

        🔴 这条不是形式主义：Task 46 首轮之所以能把空分母写成通过，正是因为没有任何
        字段迫使作者把「宣称 / 不宣称」写下来。
        """
        block = manifest_slice["property_denominators"]
        assert block["property_20"]["not_claimed_passing"] is True, (
            "slice 的 property_20.not_claimed_passing 必须为 true —— 本 slice contract 数为 0"
        )
        assert block["property_20"].get("not_claimed_passing_part"), (
            "Property 20 必须写明哪部分不宣称通过"
        )
        assert block["property_20"].get("carried_by"), (
            "Property 20 必须指出真分母承载者（G7 pilot 契约 + registry 守卫）"
        )
        for ref in block["property_20"]["carried_by"]:
            assert (ROOT / ref).exists(), f"property_20.carried_by 指向不存在的路径 {ref!r}"
        assert block["property_21"]["not_claimed_passing"] is True
        assert block["property_28"]["not_claimed_passing"] is False, (
            "Property 28 在本 slice 有真分母（模板 digest），必须宣称通过而不是回避"
        )
        assert block["property_69"].get("not_claimed_passing_part"), (
            "Property 69 的正向逐 scenario 闭合缺分母，必须写明哪部分不宣称通过"
        )
        assert block["property_70"]["not_claimed_passing"] is False


# ════════════════════════════════════════════════════════════════════════════
# Property 28：immutable definition 漂移 fail closed（真分母，现算）
# ════════════════════════════════════════════════════════════════════════════
class TestProperty28DefinitionDriftFailClosed:
    """
    **Validates: Requirements 6.10**

    本 slice 的真分母：15 个权威模板的 size + sha256、13 张 owner 模板的
    `normalized_structure_hash`（由 excel_instrumentation 现算）、17 组 sheet!cell 真读、
    G1 的 18 个 tab 名与兜底标签表逐条比对。任一字节漂移即红 —— 不是「无可漂移 identity
    所以通过」。
    """

    def test_authoritative_templates_digests_recompute(self, manifest_slice: dict) -> None:
        """每个登记模板的 size + sha256 都现算比对（Requirement 6.10 fail closed）。"""
        block = manifest_slice["authoritative_templates"]
        root = ROOT / block["root"]
        assert root.is_dir(), f"{root} 不存在"
        assert block["lock_file_policy"], "必须登记 `~$` 锁文件策略"
        assert block["reference_copy_status"] in (
            "absent_from_working_tree",
            "present_and_compared",
        )
        files = block["files"]
        on_disk = {
            p.name for p in root.iterdir() if p.is_file() and not p.name.startswith("~$")
        }
        registered = {f["name"] for f in files}
        assert registered == on_disk, (
            "登记模板集合与权威目录实况不符\n"
            f"  只在登记里: {sorted(registered - on_disk)}\n"
            f"  只在磁盘上: {sorted(on_disk - registered)}"
        )
        assert len(files) == 15, f"G 权威目录实测 15 个文件，登记 {len(files)} 个"
        for record in files:
            path = root / record["name"]
            assert path.is_file(), f"{path} 不存在"
            assert path.stat().st_size == record["size"], (
                f"{record['name']} size 漂移: 实测 {path.stat().st_size}，登记 {record['size']}"
            )
            actual = _sha256_of(path)
            assert actual == record["sha256"], (
                f"{record['name']} sha256 漂移: 实测 {actual}，登记 {record['sha256']}"
            )

    def test_normalized_structure_hash_recomputes(self, manifest_slice: dict) -> None:
        """13 张 owner 模板的 normalized_structure_hash 由生产实现现算复核。

        用的是生产代码里的同一个函数（`excel_instrumentation.normalized_structure_hash`），
        不在这里另写一份 —— 另写一份就是第二个真源，两侧漂移时谁都不红。
        """
        sys.path.insert(0, str(BACKEND))
        from app.services.workpaper_sync.excel_instrumentation import (  # noqa: PLC0415
            normalized_structure_hash,
        )

        block = manifest_slice["authoritative_templates"]
        root = ROOT / block["root"]
        checked = 0
        for record in block["files"]:
            expected = record.get("normalized_structure_hash")
            if expected is None:
                continue
            path = root / record["name"]
            actual = normalized_structure_hash(path.read_bytes())
            assert actual == expected, (
                f"{record['name']} 的 normalized_structure_hash 漂移: "
                f"实测 {actual}，登记 {expected}"
            )
            checked += 1
        assert checked == 13, (
            f"只有 {checked} 张模板登记了 normalized_structure_hash，应为 13"
            "（11 张单 entry owner + G4/G6 两张多 entry owner）⇒ Property 28 的结构分母缩水"
        )

    def test_digests_are_real_digests(self, manifest_slice: dict) -> None:
        """登记的 digest 必须是真 64 位十六进制且非全零（Requirement 2.3 同族）。"""
        digest_re = re.compile(r"^[0-9a-f]{64}$")
        for record in manifest_slice["authoritative_templates"]["files"]:
            for key in ("sha256", "normalized_structure_hash"):
                value = record.get(key)
                if value is None:
                    continue
                assert digest_re.match(value), f"{record['name']}.{key} 非法: {value!r}"
                assert set(value) != {"0"}, f"{record['name']}.{key} 是全零 hash"

    def test_template_owner_mapping_covers_every_entry_exactly_once(
        self, manifest_slice: dict
    ) -> None:
        """SR-8 的 G 循环形态：owner 可以是 1:N，但覆盖必须「每条 entry 恰一次」。

        🔴 FD-3：F 循环的守卫断言 belongs_to_entry 在 entry 上单射且满射。G 循环 13 张 owner
        模板要覆盖 17 条 entry（G4/G6 各一本册子服务 3 条），单射不成立。改成：
        `belongs_to_entry`（单值）与 `belongs_to_entries`（数组）的并集 == entry 全集，
        且每条 entry 恰被认领一次；两个字段不得同时非空。
        """
        entry_ids = {e["entry_id"] for e in manifest_slice["independent_entries"]}
        claims: list[str] = []
        multi = 0
        for record in manifest_slice["authoritative_templates"]["files"]:
            single = record.get("belongs_to_entry")
            many = record.get("belongs_to_entries")
            assert not (single and many), (
                f"{record['name']} 同时给了 belongs_to_entry 与 belongs_to_entries ⇒ 归属二义"
            )
            if single:
                assert single in entry_ids, f"{record['name']} 归属到不存在的 entry {single!r}"
                claims.append(single)
            elif many:
                assert isinstance(many, list) and len(many) >= 2, (
                    f"{record['name']} 的 belongs_to_entries 应是 ≥2 条的列表"
                )
                for eid in many:
                    assert eid in entry_ids, f"{record['name']} 归属到不存在的 entry {eid!r}"
                claims.extend(many)
                multi += 1
                assert record.get("excluded_reason"), (
                    f"{record['name']} 是 1:N 归属，必须写明为何 belongs_to_entry 记 null"
                )
            else:
                assert record.get("excluded_reason"), (
                    f"{record['name']} 没有归属 entry 却也没写 excluded_reason"
                )
        assert len(claims) == len(set(claims)), (
            f"有 entry 被多张模板认领: {sorted({c for c in claims if claims.count(c) > 1})}"
        )
        assert set(claims) == entry_ids, (
            f"每条 entry 必须恰有一张归属模板，缺 {sorted(entry_ids - set(claims))}"
        )
        declared = manifest_slice["honest_adjudication_summary"][
            "templates_serving_multiple_entries"
        ]
        assert declared == multi == 2, (
            f"服务多 entry 的模板实测 {multi} 张（G4/G6），slice 登记 {declared} 张"
        )

    def test_in_runtime_index_flag_recomputes_and_g_cycle_has_no_unreachable_workbook(
        self, manifest_slice: dict
    ) -> None:
        """登记的 in_runtime_index 必须与 _index.json 现算一致，且 G 目录无不可达冗余册。

        🔴 双向锁：D4 `D4收入底稿.xlsx` / F2 `F2存货.xlsx` 那种「磁盘有、索引无」的对象在 G
        目录**不存在**（15/15 全在索引里）。将来谁扔一本进来就打红。
        """
        indexed = _indexed_relpaths()
        root_rel = manifest_slice["authoritative_templates"]["root"].split("/")[-1]
        unreachable: list[str] = []
        for record in manifest_slice["authoritative_templates"]["files"]:
            actual = f"{root_rel}/{record['name']}" in indexed
            assert record.get("in_runtime_index") == actual, (
                f"{record['name']} 的 in_runtime_index 登记为 "
                f"{record.get('in_runtime_index')}，现算为 {actual}"
            )
            if not actual:
                unreachable.append(record["name"])
        assert unreachable == [], (
            f"G 权威目录出现未索引文件 {sorted(unreachable)} ⇒ 与冻结结论"
            "（G 循环无 D4/F2 同型不可达冗余册）不符，请更新 slice 的 "
            "template_resolution_audit.unreachable_redundant_workbook"
        )
        audit = manifest_slice["template_resolution_audit"]["unreachable_redundant_workbook"]
        assert audit["verdict"] == "not_present_in_g_cycle"

    def test_every_entry_template_ref_is_registered_with_digest(
        self, manifest_slice: dict
    ) -> None:
        """template_ref **必须非空**，且它指向的文件在 authoritative_templates 里带 digest。

        🔴 不许写成 `if ref:` —— template_ref 为 null 时整条跳过是 Task 46 首轮实测过的
        fail-open（置 null 后 36 例全绿）。
        """
        by_name = {
            f["name"]: f for f in manifest_slice["authoritative_templates"]["files"]
        }
        root = ROOT / manifest_slice["authoritative_templates"]["root"]
        for entry in manifest_slice["independent_entries"]:
            ref = entry.get("template_ref")
            assert isinstance(ref, str) and ref.strip(), (
                f"{entry['entry_id']} 的 template_ref 为空 —— 不允许缺省跳过"
            )
            name = ref.split("/")[-1]
            assert name in by_name, (
                f"{entry['entry_id']} 的 template_ref {ref!r} 未登记进 authoritative_templates"
            )
            record = by_name[name]
            owners = (
                [record["belongs_to_entry"]]
                if record.get("belongs_to_entry")
                else record.get("belongs_to_entries", [])
            )
            assert entry["entry_id"] in owners, (
                f"{entry['entry_id']} 的 template_ref 归属登记为 {owners!r}"
            )
            assert record.get("normalized_structure_hash"), (
                f"{entry['entry_id']} 的 template_ref 没有 normalized_structure_hash ⇒ "
                "结构漂移无从检出"
            )
            path = root / name
            assert path.is_file()
            assert _sha256_of(path) == record["sha256"]
            assert path.stat().st_size == record["size"]

    def test_bp5_g1_fallback_sheet_labels_point_at_nonexistent_tabs(
        self, manifest_slice: dict
    ) -> None:
        """BP-5 双向锁：G1 兜底标签表**真的**有 5 条不是权威模板 tab（不是猜的）。

        🔴 这条既锁缺陷存在，也锁修复后必须回来改 slice：把 5 条改对之后它会打红，
        逼作者把 BP-5 的 status 从 REGISTERED_NOT_FIXED 改掉。
        """
        bp5 = next(bp for bp in manifest_slice["blocking_preconditions"] if bp["id"] == "BP-5")
        assert bp5["status"] == "REGISTERED_NOT_FIXED"
        real = _sheet_names(TEMPLATE_DIR / "G" / "G1 交易性金融资产.xlsx")
        assert len(real) == 18, f"G1 权威模板实测 {len(real)} 个 tab，冻结结论是 18"
        source = _strip_ts_comments(G1_SHEET_LABELS.read_text(encoding="utf-8"))
        block = re.search(
            r"G1_SHEET_LABEL_MAP:\s*Record<string,\s*string>\s*=\s*\{([\s\S]*?)\n\}", source
        )
        assert block, "找不到 G1_SHEET_LABEL_MAP 的声明 ⇒ 判据无分母"
        pairs = re.findall(r"(?m)^\s*'?([^':\s]+)'?:\s*'([^']*)',", block.group(1))
        assert len(pairs) == 18, f"兜底标签表解析出 {len(pairs)} 条，应为 18 ⇒ 正则失配"
        mismatched = {code: label for code, label in pairs if label not in real}
        assert set(mismatched) == {"G1A", "G1-8", "G1-10", "G1-12", "附注国企"}, (
            f"兜底标签与权威 tab 不符的集合实测为 {sorted(mismatched)}，与 BP-5 冻结结论不符。"
            "若已修好，请更新 BP-5 的 status 与本判据"
        )
        # 缺陷的触发路径必须逐字可复现（否则 BP-5 的 consequence 是散文）
        dual = _strip_ts_comments(G1_DUAL_MODE.read_text(encoding="utf-8"))
        assert "return sheetName?.value || 'G1-1'" in dual, (
            "resolveOoSheetName 的空 currentSheet 兜底路径已变 —— BP-5 的触发条件要重写"
        )
        assert "return resolveG1SheetLabel(code, availableSheets.value, sheetName?.value)" in dual
        assert "return G1_SHEET_LABEL_MAP[code] ?? fallbackSheetName ?? code" in source, (
            "resolveG1SheetLabel 的兜底表回落语句已变 —— BP-5 的触发条件要重写"
        )
        # E1 宿主永远走兜底：availableSheets 传空数组
        e1 = _strip_ts_comments(E1_HOST.read_text(encoding="utf-8"))
        assert "availableSheets: computed(() => [])" in e1, (
            "E1 宿主已不再传空 availableSheets ⇒ BP-5 的 observable_consequences 第 2 条要更新"
        )
        # 跨循环命名空间碰撞：含「调整分录」一律归一成 G1-3
        assert "if (/调整分录/.test(sheetName)) return 'G1-3'" in source, (
            "extractG1SheetCode 的「调整分录」归一已变 ⇒ BP-5 的第 3 条 observable 要更新"
        )


# ════════════════════════════════════════════════════════════════════════════
# 模板解析实测：G 循环有无 D4/F2 同型缺陷（虚报与漏报同样是缺陷）
# ════════════════════════════════════════════════════════════════════════════
class TestTemplateResolution:
    """
    **Validates: Requirements 6.10**

    任务正文要求核查 G 循环有无 D4/F2 同型的「整册码回落打开错工作簿」。实测结论是
    **不存在**，所以本类做的是**正向**核验（每个码都解析对）而不是登记一条不存在的 BP。
    """

    def test_no_wrong_workbook_fallback_in_g_cycle(self, manifest_slice: dict) -> None:
        """13 个整册码 + 8 个子码逐一实测：finder 返回的都是正确册子。"""
        sys.path.insert(0, str(BACKEND))
        from app.services.wp_template_finder import (  # noqa: PLC0415
            find_all_template_files,
            find_template_file,
            find_template_file_any,
        )

        audit = manifest_slice["template_resolution_audit"]["wrong_workbook_fallback"]
        assert audit["verdict"] == "not_present_in_g_cycle"

        whole = {
            "G1": "G1 交易性金融资产.xlsx",
            "G2": "G2 应收利息.xlsx",
            "G3": "G3 应收股利.xlsx",
            "G4": "G4 债权投资.xlsx",
            "G5": "G5 长期应收款.xlsx",
            "G6": "G6 其他债权投资.xlsx",
            "G8": "G8 其他权益工具投资.xlsx",
            "G9": "G9 其他非流动金融资产.xlsx",
            "G10": "G10 交易性金融负债.xlsx",
            "G11": "G11 投资收益.xlsx",
            "G12": "G12 净敞口套期收益.xlsx",
            "G13": "G13 公允价值变动收益.xlsx",
            "G14": "G14 信用减值损失.xlsx",
        }
        for code, expected in whole.items():
            got = find_template_file(code)
            assert got is not None and got.name == expected, (
                f"find_template_file({code!r}) 返回 {got.name if got else None!r}，应为 {expected!r} "
                "⇒ G 循环出现了 D4/F2 同型的回落错册，请登记为 BP 并更新 "
                "template_resolution_audit"
            )
            got_any = find_template_file_any(code)
            assert got_any is not None and got_any.name == expected
            all_names = {p.name for p in find_all_template_files(code)}
            assert all_names == {expected}, (
                f"find_all_template_files({code!r}) 返回 {sorted(all_names)}，应只有 {expected!r} "
                "⇒ 出现了同码多册竞争（F2 那种形态）"
            )

        # 子码必须回落到父册（本 slice 每条 entry 的 primary sheet 子码各试一次）
        for entry in manifest_slice["independent_entries"]:
            sheet_code = re.search(r"(G\d+(?:-\d+)?)\s*$", entry["html_counterpart"]["primary_table"]["template_sheet"])
            assert sheet_code, (
                f"{entry['entry_id']} 的 template_sheet 尾部提不出子码 ⇒ 判据无法构造"
            )
            code = sheet_code.group(1)
            expected = entry["template_ref"].split("/")[-1]
            got = find_template_file_any(code)
            assert got is not None and got.name == expected, (
                f"find_template_file_any({code!r}) 返回 {got.name if got else None!r}，"
                f"应为 {expected!r}（{entry['entry_id']} 的 template_ref）"
            )


# ════════════════════════════════════════════════════════════════════════════
# 附注同步链路完整性：自造件处置 + metadata sheet 泄漏
# ════════════════════════════════════════════════════════════════════════════
class TestNoteSyncPathIntegrity:
    """
    **Validates: Requirements 9.1**

    任务正文两条硬约束：「先修/排除自造 G6TabDisclosureListed」与「不把虚构表或 metadata
    sheet 同步到附注」。

    🔴 关于 G6：任务正文给的二选一（(a) 重写 / (b) 排除+守卫）**前提都是「自造件仍在」**，
    而实测该前提已不成立 —— `.kiro/specs/disclosure-sync-path-buildout/` Task 2.3 已按权威
    模板重写。所以本类做的是**第三种处置：核验 + 回归锁**，理由逐条写在 slice 的
    `note_sync_path_audit.g6_listed_disclosure_forgery`（含为什么 (a)(b) 都不该做）。
    """

    def test_g6_forgery_markers_exist_only_in_comments(self, manifest_slice: dict) -> None:
        """自造标记（generateRows / 成本项目）剥掉注释后必须一处不剩。

        🔴 判据方向要对：不是「文件里不含这些字」（那会因为 docstring 里的历史说明而假红），
        而是「**可执行代码/模板**里不含」。`_strip_ts_comments` 的两条反向自检（不剥字符串
        字面量 / 真剥注释）保证这条不恒真。
        """
        audit = manifest_slice["note_sync_path_audit"]["g6_listed_disclosure_forgery"]
        assert audit["verdict"] == "already_rewritten_by_another_spec_verified_here"
        assert audit["disposal_chosen"] == "verify_and_regression_lock"
        for key in ("why_not_option_a_rewrite", "why_not_option_b_exclude", "regression_lock"):
            assert len(audit.get(key, "")) > 60, f"note_sync_path_audit 缺 {key} 的实质说明"

        forgery = re.compile(r"generateRows|成本项目")
        checked = 0
        for path in (G6_LISTED_TAB, G6_LISTED_ROWS, G6_LISTED_COMPOSABLE):
            assert path.is_file(), f"{path} 不存在"
            raw = path.read_text(encoding="utf-8")
            assert forgery.search(raw), (
                f"{path.name} 的注释里已不再记录自造历史 ⇒ 本判据失去分母"
                "（`_strip_ts_comments` 若剥过头会让下面那条恒真，靠这条兜底）"
            )
            code_only = _strip_ts_comments(raw)
            hits = sorted({m.group(0) for m in forgery.finditer(code_only)})
            assert not hits, (
                f"{path.name} 的可执行代码/模板里出现自造标记 {hits} ⇒ 自造形态回退，"
                "接附注同步会把假数据推进附注"
            )
            checked += 1
        assert checked == 3, f"只核验了 {checked} 个模块 ⇒ 分母可疑"

    def test_g6_listed_subtables_come_from_the_note_section_map(self) -> None:
        """listed 侧 14 张子表名必须由 g6NoteSectionMap 声明，且不含虚构名。"""
        source = _strip_ts_comments(G6_NOTE_SECTION_MAP.read_text(encoding="utf-8"))
        base = re.search(r"G6_LISTED_SUBTABLE\s*=\s*\{([\s\S]*?)\}\s*as const", source)
        stage = re.search(
            r"G6_LISTED_STAGE_SUBTABLE:\s*readonly\s+string\[\]\s*=\s*\[([\s\S]*?)\]", source
        )
        assert base and stage, "找不到 G6_LISTED_SUBTABLE / G6_LISTED_STAGE_SUBTABLE 声明"
        base_names = re.findall(r"'([^']+)'", base.group(1))
        stage_names = re.findall(r"'([^']+)'", stage.group(1))
        assert len(base_names) == 8, f"G6_LISTED_SUBTABLE 实测 {len(base_names)} 条，应为 8"
        assert len(stage_names) == 6, f"G6_LISTED_STAGE_SUBTABLE 实测 {len(stage_names)} 条，应为 6"
        all_names = base_names + stage_names
        assert len(all_names) == 14, "上市侧应为 14 张子表（8 + 6）"
        for name in all_names:
            for bad in ("成本项目", "其他披露事项", "其他债权投资成本"):
                assert bad not in name, f"子表名 {name!r} 含虚构小节标记 {bad!r}"
        # 组件必须真消费这两个常量（不是另写一份）
        tab = _strip_ts_comments(G6_LISTED_TAB.read_text(encoding="utf-8"))
        assert "G6_LISTED_SUBTABLE" in tab, (
            "G6TabDisclosureListed 没有 import G6_LISTED_SUBTABLE ⇒ 表名有第二份真源"
        )
        # 国企侧只 2 张（宁缺勿造），必须显式声明
        soe = re.search(r"G6_SOE_SUBTABLE\s*=\s*\{([\s\S]*?)\}\s*as const", source)
        assert soe, "找不到 G6_SOE_SUBTABLE 声明"
        assert len(re.findall(r"'([^']+)'", soe.group(1))) == 2, (
            "国企侧应只 2 张子表（三阶段系列表按宁缺勿造不推）"
        )

    def test_g6_listed_tab_is_really_in_the_note_sync_path(self) -> None:
        """反向：该组件确实**在**同步链路里（不是死代码，也不是被排除掉）。"""
        source = _strip_ts_comments(G6_LISTED_TAB.read_text(encoding="utf-8"))
        template = _vue_template(source)
        assert "syncToDisclosureNotes" in template, (
            "同步入口没进模板 ⇒ 组件在附注链路里是死代码"
        )
        assert "disclosure-notes/sync-from-workpaper" in source, (
            "找不到 disclosure-notes 同步端点调用 ⇒ 该组件并未真正接入附注链路"
        )
        assert "G6_NOTE_SECTION" in source, "同步目标章节号没有取自 g6NoteSectionMap"

    def test_no_metadata_sheet_reaches_the_note_mapping(self, manifest_slice: dict) -> None:
        """13 个 G 循环 NoteSectionMap 的披露 sheet 名必须是权威模板真实 tab，且非 metadata。"""
        audit = manifest_slice["note_sync_path_audit"]["metadata_sheet_leakage"]
        assert audit["verdict"] == "none"
        assert set(audit["forbidden_metadata_sheets"]) == set(FORBIDDEN_NOTE_SHEETS)

        sheets_by_code: dict[str, list[str]] = {}
        for path in (TEMPLATE_DIR / "G").glob("*.xlsx"):
            if path.name.startswith("~$"):
                continue
            sheets_by_code[path.name.split(" ")[0]] = _sheet_names(path)
        assert len(sheets_by_code) == 15, f"读到 {len(sheets_by_code)} 本 G 模板，应为 15"

        checked = 0
        for path in sorted(COMPOSABLES.glob("g*NoteSectionMap.ts")):
            match = re.match(r"g(\d+)NoteSectionMap\.ts", path.name)
            assert match, f"意外的文件名 {path.name}"
            code = f"G{match.group(1)}"
            if code == "G7":
                continue  # G7 已排除
            source = _strip_ts_comments(path.read_text(encoding="utf-8"))
            block = re.search(
                r"_DISCLOSURE_SHEET_NAME\s*=\s*\{([\s\S]*?)\}\s*as const", source
            )
            assert block, f"{path.name} 找不到 *_DISCLOSURE_SHEET_NAME 声明"
            names = [
                n
                for n in re.findall(r"'([^']+)'", block.group(1))
                if n not in ("listed", "soe")
            ]
            assert len(names) == 2, f"{path.name} 应声明上市/国企两个 sheet 名，实测 {names}"
            real = sheets_by_code.get(code)
            assert real, f"找不到 {code} 的权威模板"
            for name in names:
                assert name in real, (
                    f"{path.name} 声明的披露 sheet {name!r} 不是 {code} 权威模板的真实 tab；"
                    f"实有 {real}"
                )
                assert name not in FORBIDDEN_NOTE_SHEETS, (
                    f"{path.name} 把 metadata sheet {name!r} 声明成披露页 ⇒ 会同步进附注"
                )
            checked += 1
        assert checked == 13, f"只核验了 {checked} 个 NoteSectionMap，应为 13（G1~G14 除 G7）"

    def test_note_sync_registry_has_no_metadata_sheet_for_g_cycle(self) -> None:
        """registry 侧同样不得出现 metadata sheet（分母：14 条 G 条目）。"""
        registry = _load(NOTE_SYNC_REGISTRY)
        found = 0
        for item in registry["entries"]:
            wp = str(item.get("wp_code") or "")
            if not re.match(r"^G\d", wp):
                continue
            found += 1
            for key in ("workpaper_sheet", "sheet_name", "sheet"):
                value = item.get(key)
                assert value not in FORBIDDEN_NOTE_SHEETS, (
                    f"registry 里 {wp} 的 {key}={value!r} 是 metadata sheet"
                )
        assert found == 14, f"registry 里的 G 循环条目实测 {found} 条，冻结结论是 14 ⇒ 分母可疑"

    def test_g4_not_synced_tables_are_declared_instead_of_fabricated(self) -> None:
        """正向证据：源模板真有但底稿字段不支持的表，必须显式登记不推而不是造字段硬推。

        🔴 本条曾用 `re.findall(r"'([^']{20,})'", body)` 数「原因说明」，实测只捞到 5 条并据此
        判「正则失配或登记缩水」。逐对解析后事实是：**登记 6 条、原因 6 条**，漏的那条是
        `'期末重要的债权投资（续：上年年末余额）': '同上（续表）'` —— 续表的理由天然短（6 字），
        `{20,}` 的长度阈值把它筛掉了。把「长度」当「有实质说明」的代理是判据缺陷：
        续表条目正确的判法是**能追溯到本表的实质理由**，不是自己再抄一遍列差异。
        改成逐 key→reason 结构化解析后两侧都咬：少登记一条即红，把实质理由换成「同上」但
        本表没登记也红。
        """
        source = _strip_ts_comments(G4_NOTE_SECTION_MAP.read_text(encoding="utf-8"))
        block = re.search(
            r"G4_NOT_SYNCED_TABLES:\s*Readonly<Record<string,\s*string>>\s*=\s*\{([\s\S]*?)\n\}",
            source,
        )
        assert block, "找不到 G4_NOT_SYNCED_TABLES 声明 ⇒「宁缺勿造」这条正向证据不存在"
        pairs = _parse_string_record(block.group(1))
        assert len(pairs) == 6, (
            f"G4_NOT_SYNCED_TABLES 逐对解析出 {len(pairs)} 条登记（应为 6）："
            f"{[k for k, _ in pairs]} ⇒ 登记缩水或解析器失配"
        )
        registered = dict(pairs)
        assert len(registered) == len(pairs), f"表名重复: {[k for k, _ in pairs]}"
        substantive = {
            key
            for key, reason in pairs
            if "源模板" in reason and ("底稿" in reason or "无这些字段" in reason)
        }
        assert len(substantive) == 5, (
            f"实质说明列差异的登记实测 {len(substantive)} 条（应为 5）：{sorted(substantive)}"
        )
        for key, reason in pairs:
            if key in substantive:
                continue
            # 唯一合法的非实质理由 = 续表引用，且其本表必须已带实质理由
            assert "同上" in reason, (
                f"不推理由 {reason[:40]!r} 既没说明源模板与底稿的列差异，也不是续表引用 "
                "⇒ 等于自由文本豁免"
            )
            base = re.sub(r"（续[:：][^）]*）$", "", key)
            assert base != key, (
                f"{key!r} 的理由是「同上」但表名不是续表形态 ⇒ 追溯不到本表理由"
            )
            assert base in substantive, (
                f"续表 {key!r} 的理由是「同上」，但本表 {base!r} 没有带实质列差异说明的登记 "
                "⇒「同上」指向空处"
            )


# ════════════════════════════════════════════════════════════════════════════
# Property 69：evidence 与计数 —— 负向分母真验，正向逐 scenario 闭合不宣称通过
# ════════════════════════════════════════════════════════════════════════════
class TestProperty69EvidenceAndCounters:
    """
    **Validates: Requirements 12.10, 12.11**

    本 slice 没有任何 sync_test_run，故 evidence 的**正向**闭合（逐 scenario 实体 +
    服务端重算 summary）在这里没有分母 —— 那部分由 Task 44 同型的 OO 9.4 gate 承载，
    已登记 BP-4，**本类不宣称它通过**。本类真验的是负向分母与计数的可复算性。
    """

    def test_slice_scope_is_recomputable_from_the_manifest(
        self, manifest_slice: dict, full_manifest: dict
    ) -> None:
        """slice_scope.selection_rule 必须能现算出同一个 entry 集合（不是手抄清单）。"""
        scope = manifest_slice["slice_scope"]
        assert scope["cycle"] == "G"
        assert scope["document_type"] == "xlsx"
        assert scope["selection_rule"], "selection_rule 为空 ⇒ 「17」在文档里无推导"

        independent: set[str] = set()
        non_independent: set[str] = set()
        for entry in full_manifest["entries"]:
            if entry.get("document_type") != "xlsx":
                continue
            patterns = entry.get("wp_match", {}).get("wp_code_patterns", [])
            if not any(str(p).startswith("G") for p in patterns):
                continue
            if entry.get("independent_entry"):
                independent.add(entry["entry_id"])
            else:
                non_independent.add(entry["entry_id"])
        assert len(independent) == scope["g_prefixed_independent_total"], (
            f"现算 G 码 xlsx independent={len(independent)}，slice 写 "
            f"{scope['g_prefixed_independent_total']}"
        )
        assert len(non_independent) == scope["non_independent_entry_count"], (
            f"现算 non_independent={len(non_independent)}，slice 写 "
            f"{scope['non_independent_entry_count']}"
        )
        assert non_independent == {UNREACHABLE_STUB_ENTRY_ID}, (
            f"G 循环的非独立 entry 实测为 {sorted(non_independent)}，冻结结论是那一个不可达旧桩"
        )
        assert independent & G7_ENTRY_IDS == G7_ENTRY_IDS, (
            "G7 的三条 entry 应在 manifest 的 independent 集合里（本 slice 主动排除它们）"
        )
        expected = independent - G7_ENTRY_IDS
        assert len(expected) == scope["independent_entry_count"], (
            f"现算 slice entry 数={len(expected)}，slice 写 {scope['independent_entry_count']}"
        )
        assert expected == {e["entry_id"] for e in manifest_slice["independent_entries"]}, (
            "slice 的 entry 集合与按 selection_rule 现算的集合不等"
        )
        assert scope["excluded_g7_entry_count"] == len(G7_ENTRY_IDS)

    def test_no_g_docx_entry_exists_in_the_source_manifest(self, full_manifest: dict) -> None:
        """G 循环无 Word 通道 entry —— 若哪天出现，本条打红提醒重算 selection_rule。"""
        docx_entries = [e for e in full_manifest["entries"] if e.get("document_type") == "docx"]
        assert docx_entries, "全量 manifest 里一条 docx entry 都没有 ⇒ 判据无分母"
        g_docx = [
            e["entry_id"]
            for e in docx_entries
            if any(
                str(p).startswith("G")
                for p in e.get("wp_match", {}).get("wp_code_patterns", [])
            )
        ]
        assert not g_docx, (
            f"source manifest 里出现 G 码 docx entry {g_docx} ⇒ selection_rule 要加 Word 通道分支"
        )

    def test_excluded_items_have_reasons(self, manifest_slice: dict) -> None:
        excluded = manifest_slice["slice_scope"]["excluded_from_slice"]
        assert len(excluded) >= 4, (
            f"excluded_from_slice 只有 {len(excluded)} 条 —— G7 三条 entry / G0 函证族 / "
            "不可达旧桩 / 附注同步链路四类都要显式排除"
        )
        for item in excluded:
            assert item.get("what")
            assert len(item.get("reason", "")) > 20
        # G7 的排除必须逐条列出 entry_id（只写「G7」无法复核）
        g7_text = " ".join(item["what"] + item["reason"] for item in excluded)
        for eid in G7_ENTRY_IDS:
            assert eid in g7_text, f"excluded_from_slice 没有显式列出 G7 的 entry {eid!r}"

    def test_each_entry_has_unique_id(self, manifest_slice: dict) -> None:
        ids = [e["entry_id"] for e in manifest_slice["independent_entries"]]
        assert len(ids) == len(set(ids))

    def test_evidence_state_and_reasons(self, manifest_slice: dict) -> None:
        """SR-7：UNVERIFIABLE 必须带非空 unverifiable_reasons；字段不得缺失。"""
        for entry in manifest_slice["independent_entries"]:
            evidence = entry.get("evidence")
            assert isinstance(evidence, dict), (
                f"{entry['entry_id']} 的 evidence 缺失 —— 连 UNVERIFIABLE 都没标的条目"
                "在 stale 统计里无法归类（AC 12.13）"
            )
            for key in (
                "browser_case",
                "contract_test",
                "sync_test_run_id",
                "required_scenario_set_digest",
                "verification_state",
            ):
                assert key in evidence, f"{entry['entry_id']} 的 evidence 缺 {key}"
            state = evidence["verification_state"]
            assert state in ("UNVERIFIABLE", "VERIFIED", "FAILED"), (
                f"{entry['entry_id']} 的 verification_state={state!r} 非法"
            )
            if state == "UNVERIFIABLE":
                reasons = evidence.get("unverifiable_reasons")
                assert isinstance(reasons, list) and reasons, (
                    f"{entry['entry_id']} 标 UNVERIFIABLE 却没说明原因 = 自由文本豁免"
                )

    def test_contract_test_points_at_this_file(self, manifest_slice: dict) -> None:
        """evidence.contract_test 必须指向真存在的守卫文件（且就是本文件）。"""
        expected = "backend/tests/workpaper_sync/test_task49_g_cycle_migration.py"
        for entry in manifest_slice["independent_entries"]:
            declared = entry["evidence"]["contract_test"]
            assert declared == expected, (
                f"{entry['entry_id']} 的 contract_test={declared!r}，应为 {expected!r}"
            )
            assert (ROOT / declared).is_file()

    def test_no_entry_claims_verified_without_a_test_run(self, manifest_slice: dict) -> None:
        """假双向已验收计数的真分母：VERIFIED 必须有 test run + scenario digest。"""
        for entry in manifest_slice["independent_entries"]:
            evidence = entry["evidence"]
            if evidence["verification_state"] == "VERIFIED":
                assert evidence["sync_test_run_id"], (
                    f"{entry['entry_id']} 声称 VERIFIED 但没有 sync_test_run_id"
                )
                assert evidence["required_scenario_set_digest"], (
                    f"{entry['entry_id']} 声称 VERIFIED 但没有 required_scenario_set_digest"
                )

    def test_summary_counters_are_recomputed_from_entries(self, manifest_slice: dict) -> None:
        """SR-1 / SR-2：摘要计数逐项由 entries 现算，抄错或凑数即红。"""
        entries = manifest_slice["independent_entries"]
        summary = manifest_slice["honest_adjudication_summary"]
        assert summary["total_independent"] == len(entries)
        assert summary["total_independent"] == manifest_slice["slice_scope"][
            "independent_entry_count"
        ]
        for capability, key in (
            ("bidirectional", "adjudicated_as_bidirectional"),
            ("single_onlyoffice", "adjudicated_as_single_onlyoffice"),
            ("single_html", "adjudicated_as_single_html"),
            ("unreachable", "adjudicated_as_unreachable"),
        ):
            expected = sum(1 for e in entries if e.get("capability") == capability)
            assert summary[key] == expected, f"{key} 应为 {expected}，slice 写 {summary[key]}"
        pending = sum(1 for e in entries if e.get("capability") is None)
        assert summary["capability_verdict_pending"] == pending
        assert summary["slice_counters"]["unadjudicated"] == pending, (
            "capability 终态未定的条目必须计入 unadjudicated —— 归零不是进度"
        )
        assert summary["html_counterpart_exists"] == sum(
            1 for e in entries if e.get("html_counterpart_verdict") == "exists"
        )
        assert summary["html_counterpart_none"] == sum(
            1 for e in entries if e.get("html_counterpart_verdict") == "none"
        )
        assert summary["entries_left_unverifiable"] == sum(
            1 for e in entries if e["evidence"]["verification_state"] == "UNVERIFIABLE"
        )
        assert summary["entries_with_positional_row_identity_defect"] == sum(
            1 for e in entries if e["html_counterpart"].get("row_identity_is_positional")
        )
        assert summary["slice_counters"]["fake_bidirectional_claimed_verified"] == sum(
            1
            for e in entries
            if e.get("capability") == "bidirectional"
            and e["evidence"]["verification_state"] == "VERIFIED"
            and e.get("adapter_id") is None
        )
        assert summary["slice_counters"]["bidirectional_unverified"] == sum(
            1
            for e in entries
            if e.get("capability") == "bidirectional"
            and e["evidence"]["verification_state"] != "VERIFIED"
        )
        assert summary["pilot_contract_published"] == 0, (
            "本 slice 没有 pilot（G 循环的 pilot 是已排除的 G7）⇒ 该计数必须是 0"
        )
        assert summary["finalized_published_representations"] == 0

    def test_form_difference_counts_are_recomputed(self, manifest_slice: dict) -> None:
        """四组形态分布字典必须逐键由 entries 现算（凑数即红）。"""
        entries = manifest_slice["independent_entries"]
        summary = manifest_slice["honest_adjudication_summary"]
        for counter_key, field in (
            ("transport_key_kind_counts", "transport_key_kind"),
            ("payload_column_mode_counts", "payload_column_mode"),
            ("http_client_binding_counts", "http_client_binding"),
            ("row_identity_kind_counts", "row_identity_kind"),
        ):
            actual: dict[str, int] = {}
            for entry in entries:
                value = entry["html_counterpart"][field]
                actual[value] = actual.get(value, 0) + 1
            assert summary[counter_key] == actual, (
                f"{counter_key} 现算 {actual}，slice 写 {summary[counter_key]}"
            )
            assert sum(actual.values()) == len(entries), (
                f"{counter_key} 的计数之和 {sum(actual.values())} != entry 数 {len(entries)}"
            )

    def test_form_differences_block_is_complete(self, manifest_slice: dict) -> None:
        """g_cycle_form_differences 必须逐条带 what + 可复算来源，不是散文。

        🔴 本条曾读 `why_this_key_exists`（那是 `property_denominators` 的键名，被复制过来
        时没改），对 `g_cycle_form_differences` 恒 KeyError。键名是判据的一部分：读错键的
        守卫既不是「通过」也不是「发现缺陷」，只是判据本身坏了。
        """
        block = manifest_slice["g_cycle_form_differences"]
        assert "why_this_key_exists" not in block, (
            "g_cycle_form_differences 的说明键是 why_this_section_exists；出现 "
            "why_this_key_exists 说明两个块的键名开始互串，本判据要跟着改"
        )
        assert len(block["why_this_section_exists"]) > 60
        diffs = block["differences"]
        assert len(diffs) >= 5, f"只登记了 {len(diffs)} 条形态差异，实测 5 条"
        ids = [d["id"] for d in diffs]
        assert len(ids) == len(set(ids)), f"形态差异 id 重复: {ids}"
        for diff in diffs:
            for key in ("id", "name", "what", "counted_from"):
                assert diff.get(key), f"形态差异 {diff.get('id')} 缺 {key}"
            assert len(diff["what"]) > 40, f"{diff['id']} 的 what 过短"
            for key in ("source_of_truth",):
                refs = diff.get(key)
                if refs is None:
                    continue
                if isinstance(refs, str):
                    refs = [refs]
                for ref in refs:
                    if ref.startswith("independent_entries"):
                        continue
                    assert (ROOT / ref.split("#")[0]).exists(), (
                        f"{diff['id']} 的 {key} 指向不存在的路径 {ref!r}"
                    )

    def test_blocking_preconditions_are_complete(self, manifest_slice: dict) -> None:
        """每条阻断项必须有 id/blocks/what/后果/status/must_fix_before/owner_task/source_refs。

        「后果」在事实范式里有两种形态（单串 `consequence` / 数组 `observable_consequences`），
        按 one-of 校验：至少一个非空。
        """
        blocking = manifest_slice["blocking_preconditions"]
        assert len(blocking) >= 10, f"只登记了 {len(blocking)} 条阻断项，可疑"
        ids = [bp["id"] for bp in blocking]
        assert len(ids) == len(set(ids)), f"阻断项 id 重复: {ids}"
        for bp in blocking:
            for key in ("id", "blocks", "what", "status", "must_fix_before", "owner_task"):
                assert bp.get(key), f"阻断项 {bp.get('id')} 缺 {key}"
            assert bp.get("consequence") or bp.get("observable_consequences"), (
                f"阻断项 {bp['id']} 既没有 consequence 也没有 observable_consequences —— "
                "后果必须写明，否则它只是一句「有问题」"
            )
            assert bp.get("why_not_fixed_here"), (
                f"阻断项 {bp['id']} 没写 why_not_fixed_here ⇒ 「只登记不修」缺理由"
            )
            refs = bp.get("source_refs")
            assert isinstance(refs, list) and refs, (
                f"阻断项 {bp['id']} 没有 source_refs —— 无据可查的阻断项等于自由文本"
            )
            for ref in refs:
                raw = ref.split("#")[0].strip()
                assert (ROOT / raw).exists(), (
                    f"阻断项 {bp['id']} 的 source_ref {ref!r} 指向不存在的路径"
                )


# ════════════════════════════════════════════════════════════════════════════
# Property 70：不得跨 entry 复用
# ════════════════════════════════════════════════════════════════════════════
class TestProperty70NoCrossEntryReuse:
    """
    **Validates: Requirements 12.12**
    """

    def test_cross_entry_isolation_block_present(self, manifest_slice: dict) -> None:
        isolation = manifest_slice["cross_entry_isolation"]
        assert isolation["rule"]
        assert len(isolation["assertions"]) >= 6

    def test_g_entries_disjoint_from_d_e_f_slices(self, manifest_slice: dict) -> None:
        """D / E / F / G 四个 slice 的 entry 集合两两不相交（重复计数 = 假迁移进度）。"""
        g_ids = {e["entry_id"] for e in manifest_slice["independent_entries"]}
        for label, path in (("D", D_SLICE_PATH), ("E", E_SLICE_PATH), ("F", F_SLICE_PATH)):
            assert path.is_file(), f"{path} 不存在（{label} 循环交付物）"
            other = {e["entry_id"] for e in _load(path)["independent_entries"]}
            assert other, f"{label} slice 的 entry 集合为空 ⇒ 不相交判据恒真"
            assert not (g_ids & other), f"G/{label} slice entry 相交: {sorted(g_ids & other)}"

    def test_no_entry_carries_a_per_entry_contract_field(self, manifest_slice: dict) -> None:
        """本 slice 一条 contract 都没有 ⇒ 任何 entry 都不得挂 per_entry_contract。"""
        for entry in manifest_slice["independent_entries"]:
            assert "per_entry_contract" not in entry, (
                f"{entry['entry_id']} 挂了 per_entry_contract —— 本 slice 无契约，"
                "不得借用 G7 或其他循环的 contract_id"
            )

    def test_same_workbook_entries_do_not_share_identity(self, manifest_slice: dict) -> None:
        """同册多 entry（G4/G6 各 3 条）的身份要素必须两两互异。

        🔴 这是 G 循环特有的最大复用风险面：三条 entry 共享一本工作簿与一个 wp_code_pattern，
        若 primary table / owner / localStorage 前缀有任一相同，就等于「一条 entry 的对端结论
        被另一条借用」。
        """
        by_template: dict[str, list[dict]] = {}
        for entry in manifest_slice["independent_entries"]:
            by_template.setdefault(entry["template_ref"], []).append(entry)
        groups = {k: v for k, v in by_template.items() if len(v) > 1}
        assert len(groups) == 2, (
            f"同册多 entry 的分组实测 {len(groups)} 组（应为 G4/G6 两组）：{sorted(groups)}"
        )
        prefixes: dict[str, str] = {}
        for template, group in groups.items():
            assert len(group) == 3, f"{template} 服务 {len(group)} 条 entry，冻结结论是 3"
            for field in ("entry_id", "host_path", "legacy_dual_mode_composable"):
                values = [e[field] for e in group]
                assert len(set(values)) == len(values), (
                    f"{template} 的三条 entry 的 {field} 有重复: {values}"
                )
            items = [e["html_counterpart"]["primary_table"]["item_id"] for e in group]
            assert len(set(items)) == 3, f"{template} 的三条 entry 主表键有重复: {items}"
            owners = [e["html_counterpart"]["primary_table"]["owner_module"] for e in group]
            assert len(set(owners)) == 3, f"{template} 的三条 entry owner 模块有重复: {owners}"
            # wp_code_pattern 确实相同（这正是 BP-8 的碰撞面）
            patterns = {e["wp_code_pattern"] for e in group}
            assert len(patterns) == 1, (
                f"{template} 的三条 entry 的 wp_code_pattern 不同 {patterns} ⇒ BP-8 的前提要重写"
            )
        # localStorage 前缀全局互异（跨 17 条，不只同册三条）
        for entry in manifest_slice["independent_entries"]:
            mod = entry["legacy_dual_mode_composable"]
            path = None
            for base in (COMPOSABLES, GLOBAL_COMPOSABLES):
                candidate = base / f"{mod}.ts"
                if candidate.is_file():
                    path = candidate
                    break
            assert path is not None, f"{entry['entry_id']} 的 dual-mode 模块 {mod} 不存在"
            source = path.read_text(encoding="utf-8")
            found = re.search(r"STORAGE_PREFIX\s*=\s*'([^']*)'", source)
            if found is None:
                continue  # useG2DualMode 走共享基座，前缀由基座管
            prefix = found.group(1)
            assert prefix not in prefixes, (
                f"{entry['entry_id']} 的 localStorage 前缀 {prefix!r} 与 "
                f"{prefixes[prefix]} 相同 ⇒ 两条 entry 的模式偏好会互相顶掉"
            )
            prefixes[prefix] = entry["entry_id"]
        assert len(prefixes) == 16, (
            f"实测 {len(prefixes)} 个独立前缀，应为 16（17 条 entry 减去走共享基座的 G2）"
        )
        bp8 = next(bp for bp in manifest_slice["blocking_preconditions"] if bp["id"] == "BP-8")
        assert bp8["status"] == "REGISTERED_NOT_FIXED"

    def test_deletion_plan_entries_are_distinct(self, deletion_plan: dict) -> None:
        ids = [e["entry_id"] for e in deletion_plan["entries"]]
        assert len(ids) == len(set(ids))

    def test_deletion_plan_composables_are_distinct_and_real(self, deletion_plan: dict) -> None:
        """legacy composable 不得跨 entry 复用，且每条必须真在磁盘上。"""
        all_composables: list[str] = []
        for entry in deletion_plan["entries"]:
            for comp in entry.get("legacy_composables_to_delete", []):
                all_composables.append(comp["file"])
                assert (ROOT / comp["file"]).is_file(), (
                    f"deletion plan 里的 {comp['file']} 不存在 ⇒ 清册与磁盘脱钩"
                )
        assert len(all_composables) == len(set(all_composables)), (
            f"legacy composable 跨 entry 重复: {all_composables}"
        )
        assert len(all_composables) == 17, (
            f"清册列了 {len(all_composables)} 个 composable，应为 17（每 entry 一个）"
        )

    def test_deletion_plan_matches_slice(
        self, deletion_plan: dict, manifest_slice: dict
    ) -> None:
        """deletion plan 的裁决字段必须与 slice 一致，且理由不循环论证。"""
        by_id = {e["entry_id"]: e for e in manifest_slice["independent_entries"]}
        assert {e["entry_id"] for e in deletion_plan["entries"]} == set(by_id), (
            "deletion plan 与 slice 的 entry 集合不等 ⇒ 有 entry 漏登或多登"
        )
        for entry in deletion_plan["entries"]:
            sliced = by_id[entry["entry_id"]]
            assert entry.get("capability_adjudication") == sliced.get("capability"), (
                f"{entry['entry_id']} 的 deletion plan 裁决与 slice 不一致"
            )
            assert entry.get("html_counterpart_verdict") == sliced.get(
                "html_counterpart_verdict"
            )
            assert entry.get("capability_target") == sliced.get("capability_target")
            assert entry.get("capability_verdict_stage") == sliced.get(
                "capability_verdict_stage"
            )
            assert len(entry.get("adjudication_reason", "")) > 40
            assert entry.get("host_component") == sliced["host_path"]
            assert entry.get("ui_toolbar_gate") == sliced["ui_toolbar_gate"]
            assert entry.get("notice_mounted") is True
            comp = entry["legacy_composables_to_delete"][0]
            assert comp["exported"] == sliced["legacy_dual_mode_composable"], (
                f"{entry['entry_id']} 的 deletion plan composable 与 slice 登记不一致"
            )

    def test_deletion_plan_reasons_are_not_circular(
        self, deletion_plan: dict, paradigm: dict
    ) -> None:
        """deletion plan 的裁决理由同样不得循环论证（两份产物同一把尺）。"""
        markers = paradigm["adjudication_criteria"]["anti_patterns"][0][
            "circular_reason_markers"
        ]
        for entry in deletion_plan["entries"]:
            text = entry.get("adjudication_reason", "")
            hits = [m for m in markers if m in text]
            assert not hits, (
                f"deletion plan 里 {entry['entry_id']} 的理由命中循环论证标记 {hits}"
            )

    def test_deletion_plan_declares_no_pilot_and_preserves_shared_base(
        self, deletion_plan: dict
    ) -> None:
        """本 slice 无 pilot；共享基座不得进删除清册。

        🔴 与 F 循环不同：G 循环**有一条**包了共享基座（useG2DualMode），所以
        `g_cycle_does_not_consume_it` 必须是 false 并列出消费方 —— 写 true 就是假绿。
        """
        pilot = deletion_plan["pilot_already_migrated"]
        assert pilot["entry_id"] is None and pilot["note"], (
            "本 slice 无 pilot 必须显式写 null + 说明，不能省略字段"
        )
        preserved = deletion_plan["shared_base_preserved"]
        assert "useWorkpaperEntryDualMode.ts" in preserved["file"]
        assert preserved["g_cycle_does_not_consume_it"] is False, (
            "G 循环实测有一条 composable 包了共享基座（useG2DualMode），"
            "声明 true 会让「删基座是安全的」这个错误结论通过"
        )
        assert preserved["consumers_in_g_cycle"], "声称有消费方却没列出来"
        assert SHARED_BASE.is_file(), "共享基座必须存在"
        to_delete = {
            comp["file"]
            for entry in deletion_plan["entries"]
            for comp in entry["legacy_composables_to_delete"]
        }
        assert preserved["file"] not in to_delete, "共享基座被列入删除清册"
        for consumer in preserved["consumers_in_g_cycle"]:
            source = _strip_ts_comments((ROOT / consumer).read_text(encoding="utf-8"))
            assert "useWorkpaperEntryDualMode" in source, (
                f"{consumer} 被登记为基座消费方，但源码里没有引用 ⇒ 登记不成立"
            )

    def test_wraps_shared_base_flag_is_source_backed_both_ways(
        self, deletion_plan: dict
    ) -> None:
        """`wraps_shared_base` 双向核验：声称包了的真包，声称没包的真没包。"""
        wrapped: list[str] = []
        checked = 0
        for entry in deletion_plan["entries"]:
            for comp in entry["legacy_composables_to_delete"]:
                source = _strip_ts_comments((ROOT / comp["file"]).read_text(encoding="utf-8"))
                actual = "useWorkpaperEntryDualMode" in source
                assert actual == bool(comp.get("wraps_shared_base")), (
                    f"{comp['file']}: wraps_shared_base 声称 "
                    f"{comp.get('wraps_shared_base')}，实测 {actual}"
                )
                if actual:
                    wrapped.append(comp["file"])
                checked += 1
        assert checked == 17, f"只核验了 {checked} 个 composable ⇒ 分母可疑"
        assert len(wrapped) == 1 and wrapped[0].endswith("useG2DualMode.ts"), (
            f"包共享基座的 composable 实测为 {wrapped}，冻结结论是仅 useG2DualMode"
        )

    def test_cross_cycle_shared_composable_is_registered_with_its_consumers(
        self, deletion_plan: dict, manifest_slice: dict
    ) -> None:
        """BP-9：useG1DualMode 被 E 循环共用，必须逐字登记且消费方可复核。

        🔴 这条同时回答了 E 循环 deletion plan 的「Task 49 完成即解锁」问题：不能解锁，
        因为 Task 49 不含 step 9（宿主改线）。判据锁死两件事：① 消费方确实是两个循环的宿主；
        ② 登记里明确写了 `can_it_be_released_now: false` 与重写后的解锁条件。
        """
        block = deletion_plan["cross_cycle_shared_composable"]
        assert block["file"].endswith("useG1DualMode.ts")
        consumers = block["consumers"]
        assert len(consumers) == 2, f"登记了 {len(consumers)} 个消费方，实测应为 2"
        for consumer in consumers:
            source = _strip_ts_comments((ROOT / consumer).read_text(encoding="utf-8"))
            assert re.search(r"import\s*\{?[^}\n]*useG1DualMode", source), (
                f"{consumer} 被登记为 useG1DualMode 的消费方，但源码里没有该 import"
            )
        # 反向：全仓生产代码里的消费方恰是这两个（多一个就说明登记漏了）
        actual: set[str] = set()
        for path in FRONTEND.rglob("*"):
            if path.suffix not in (".ts", ".vue") or not path.is_file():
                continue
            rel = path.relative_to(ROOT).as_posix()
            if "__tests__" in rel or path.name == "useG1DualMode.ts":
                continue
            text = _strip_ts_comments(path.read_text(encoding="utf-8", errors="replace"))
            if re.search(r"\buseG1DualMode\b", text):
                actual.add(rel)
        assert actual == set(consumers), (
            f"useG1DualMode 的生产消费方实测 {sorted(actual)}，登记 {sorted(consumers)}"
        )
        deferred = block["e_cycle_deferred_registration"]
        assert deferred["can_it_be_released_now"] is False, (
            "E 循环的延后删除登记**不能**在本任务后解除（Task 49 不含 step 9 宿主改线）"
        )
        assert len(deferred["why_not"]) > 80
        assert len(deferred["restated_unlock_condition"]) > 40
        assert deferred["blocking_precondition"] == "BP-9"
        bp9 = next(bp for bp in manifest_slice["blocking_preconditions"] if bp["id"] == "BP-9")
        assert bp9["status"] == "REGISTERED_NOT_FIXED"

    def test_unreachable_stub_is_registered_and_still_has_zero_inbound_edges(
        self, deletion_plan: dict, manifest_slice: dict
    ) -> None:
        """AC 1.7：不可达旧桩必须登记待删，且守卫锁死它仍然零入边。

        🔴 双向锁：若有人给旧桩接了入边（把它救活），本条打红；若旧桩已被删除，
        `UNREACHABLE_STUB.is_file()` 会打红，逼作者回来更新 slice 与清册。
        """
        stub = deletion_plan["unreachable_stub_to_delete"]
        assert stub["entry_id"] == UNREACHABLE_STUB_ENTRY_ID
        assert stub["blocking_precondition"] == "BP-11"
        assert stub["inbound_reference_count"] == 0
        assert UNREACHABLE_STUB.is_file(), (
            "GtG6OtherBondEcl.vue 已不在磁盘上 ⇒ BP-11 已被处理，请更新 slice 与清册"
        )
        registry = _strip_ts_comments(HTML_RENDERER_REGISTRY.read_text(encoding="utf-8"))
        # 🔴 判「旧桩是否被救活」只能按**模块边**（import 的路径），不能按符号名。
        # 实测事实：htmlRendererRegistry#L378 有一个**同名局部别名**
        #   `const GtG6OtherBondEcl = defineAsyncComponent(() => import('./GtG6OtherBondInvestmentEcl.vue'))`
        # —— 名字叫 GtG6OtherBondEcl，指向的却是真正在用的 GtG6OtherBondInvestmentEcl.vue。
        # 按名字判会把这个别名当成「引用了旧桩」（Task 49 首轮就是这么红的，且
        # `registry.replace("GtG6OtherBondInvestmentEcl", "")` 这种先消长名再查短名的写法
        # 反而正好把别名声明行留成 `const GtG6OtherBondEcl = … import('./.vue')`）。
        # 旧桩文件的真实入边数 = 0：全仓没有任何 import 指向 `GtG6OtherBondEcl.vue`，
        # 也没有任何模板用 `<GtG6OtherBondEcl` / `<gt-g6-other-bond-ecl` 标签
        # （后者会经 unplugin-vue-components 的全局注册解析到该文件）。
        alias = stub["registry_homonymous_alias"]
        assert alias["declared_in"].endswith("htmlRendererRegistry.ts"), (
            "清册必须登记同名别名所在文件 —— 否则下一个人还会按名字判一次"
        )
        assert alias["alias_name"] == "GtG6OtherBondEcl"
        assert alias["resolves_to"] == "GtG6OtherBondInvestmentEcl.vue"
        assert alias["is_an_inbound_edge_to_the_stub"] is False
        assert len(alias["why_name_based_probes_misfire"]) > 60
        alias_decl = re.search(
            r"const\s+GtG6OtherBondEcl\s*=\s*defineAsyncComponent\(\s*\(\)\s*=>\s*"
            r"import\(\s*'\./([A-Za-z0-9_.\-]+)'\s*\)\s*\)",
            registry,
        )
        assert alias_decl, (
            "htmlRendererRegistry 里找不到登记的同名别名声明 ⇒ 别名形态已变，"
            "本判据与清册的 registry_homonymous_alias 都要重算"
        )
        assert alias_decl.group(1) == alias["resolves_to"], (
            f"同名别名实际指向 {alias_decl.group(1)!r}，清册登记 {alias['resolves_to']!r}"
        )
        assert "GtG6OtherBondInvestmentEcl" in registry, (
            "htmlRendererRegistry 里找不到真正在用的 GtG6OtherBondInvestmentEcl ⇒ 判据无分母"
        )
        # 全仓按模块边扫：没有任何生产模块 import 旧桩文件，也没有模板用它的标签。
        # 🔴 探针只认**模块说明符**（`from '…'` / `import('…')` / `require('…')`），不认
        # 「任何以该文件名结尾的字符串」—— 后者会把 workpaperSyncManifest.generated.ts 里
        # `"hostPath": "…/GtG6OtherBondEcl.vue"` 这种**数据登记**当成 import。那条登记恰恰是
        # 旧桩作为 unreachable entry 被登记的地方，把它算成入边会得出「旧桩被救活」的反结论。
        stub_import = re.compile(
            r"""(?:from|import|require)\s*\(?\s*['"][^'"]*GtG6OtherBondEcl\.vue['"]"""
        )
        stub_tag = re.compile(r"<\s*(?:GtG6OtherBondEcl|gt-g6-other-bond-ecl)[\s/>]")
        importers: list[str] = []
        scanned = 0
        for path in FRONTEND.rglob("*"):
            if path.suffix not in (".ts", ".vue") or not path.is_file():
                continue
            rel = path.relative_to(ROOT).as_posix()
            if "__tests__" in rel or path.name == "GtG6OtherBondEcl.vue":
                continue
            if rel.endswith("components.d.ts"):
                continue  # unplugin-vue-components 的生成产物，非运行时边（下面单独核）
            scanned += 1
            text = _strip_ts_comments(path.read_text(encoding="utf-8", errors="replace"))
            if stub_import.search(text) or stub_tag.search(text):
                importers.append(rel)
        assert scanned > 500, f"只扫了 {scanned} 个前端模块 ⇒ 分母可疑，入边结论不可信"
        assert not importers, (
            f"不可达旧桩被以下模块按路径 import 或按标签使用: {importers} ⇒ 它不再是旧桩，"
            "manifest 的 unreachable 裁决与 BP-11 都要重算"
        )
        # 反向自检：探针不是恒假 —— 同一把尺量真正在用的那个组件必须命中
        live_import = re.compile(
            r"""(?:from|import|require)\s*\(?\s*['"][^'"]*GtG6OtherBondInvestmentEcl\.vue['"]"""
        )
        assert live_import.search(registry), (
            "按路径扫的探针连 registry 对 GtG6OtherBondInvestmentEcl.vue 的 import 都扫不到 "
            "⇒ 探针恒假，「零入边」这个结论没有信息量"
        )
        # 生成的前端 manifest 侧：旧桩必须仍以 unreachable entry 的身份被登记（它是「数据登记」
        # 而不是入边）。少了它就说明 manifest 重新生成时把这条 AC 1.7 义务丢了。
        generated = (SYNC_DIR / "workpaperSyncManifest.generated.ts").read_text(encoding="utf-8")
        stub_block = re.search(
            r"\{[^{}]*\"entryId\":\s*\"" + re.escape(UNREACHABLE_STUB_ENTRY_ID) + r"\"[\s\S]{0,900}",
            generated,
        )
        assert stub_block, (
            f"生成的前端 manifest 里找不到 {UNREACHABLE_STUB_ENTRY_ID} ⇒ AC 1.7 的义务对象"
            "从清册里消失了"
        )
        block_text = stub_block.group(0)
        assert '"capability": "unreachable"' in block_text
        assert '"migrationState": "unreachable_pending_delete"' in block_text
        assert '"independentEntry": false' in block_text
        assert "GtG6OtherBondEcl.vue" in block_text, (
            "生成 manifest 里该 entry 的 hostPath 不再指向旧桩 ⇒ 裁决对象变了"
        )
        # 生成物侧：components.d.ts 仍会为旧桩生成全局组件声明（它按目录扫全部 .vue）。
        # 这不是运行时入边，但必须登记 —— 否则删桩时会有人以为它「被 d.ts 引用着」。
        components_dts = FRONTEND / "components.d.ts"
        assert components_dts.is_file()
        dts_declares = "GtG6OtherBondEcl.vue" in components_dts.read_text(encoding="utf-8")
        assert dts_declares == bool(stub["generated_dts_declares_it"]), (
            f"components.d.ts 是否声明旧桩：实测 {dts_declares}，清册登记 "
            f"{stub['generated_dts_declares_it']}"
        )
        bp11 = next(bp for bp in manifest_slice["blocking_preconditions"] if bp["id"] == "BP-11")
        assert bp11["status"] == "REGISTERED_NOT_FIXED"
        # 它不得混进 slice 的 independent_entries（那会把 17 变 18）
        assert UNREACHABLE_STUB_ENTRY_ID not in {
            e["entry_id"] for e in manifest_slice["independent_entries"]
        }


# ════════════════════════════════════════════════════════════════════════════
# AC 1.4：裁决自带的 UI 义务（范式 step 11）
# ════════════════════════════════════════════════════════════════════════════
class TestAc14HonestModeVisibility:
    """
    **Validates: Requirements 14.1**

    AC 1.4：未注册 adapter 的入口**禁止**显示「可双向回写」，并**必须**显示可操作原因。

    🔴 只改 slice 的 capability 而界面照旧 = AP-4 capability_without_ui_change。
    判据必须落到**模板形态**（宿主真挂了组件 + 传了自己的 entry id + 落在模式工具栏区块内），
    不能只 grep 符号名 —— 「模型声明了而模板零引用」是 G7 两级表头 0/38 的同型结构性死代码。

    AC 1.5 为什么不在这里：它的触发条件是「入口被裁决为纯 OO 或纯 HTML」。17 条 entry 都不是
    single_*，AC 1.5 的 WHEN 不成立；收掉切换按钮反而会砍掉真实可用的 OO 独立编辑通道。
    见 slice 的 BP-12。
    """

    def test_notice_module_is_the_single_source(self) -> None:
        """复用 Task 46 收口新建的单一真源，不新建第二份文案模块。"""
        source = _strip_ts_comments(NOTICE_MODULE.read_text(encoding="utf-8"))
        assert "export function entrySyncNotice" in source
        assert "SYNC_ADAPTER_REGISTERED_ENTRY_IDS" in source
        assert "ENTRY_SYNC_NOTICE_SUMMARY" in source
        assert "ENTRY_SYNC_NOTICE_REASON" in source
        siblings = sorted(p.name for p in SYNC_DIR.glob("*EntrySyncNotice*.ts"))
        assert siblings == ["workpaperEntrySyncNotice.ts"], (
            f"sync/ 下出现了多份通知文案模块 {siblings} ⇒ 单一真源被破坏"
        )

    def test_notice_component_renders_summary_and_binds_entry_id(self) -> None:
        """常显摘要 + entry 绑定必须在模板里（不是只在 script 里算）。"""
        source = _strip_ts_comments(NOTICE_COMPONENT.read_text(encoding="utf-8"))
        template = _vue_template(source)
        assert "notice.summary" in template, (
            "摘要没进模板 ⇒ 可操作原因默认不可见（EP tooltip 内容是 teleport 且 hover 才挂载）"
        )
        assert 'v-if="notice"' in template, "缺 v-if ⇒ 组件恒显，已接双向的 entry 也会挂警告"
        assert ":data-entry-sync-notice=" in template, "entry 绑定没进模板 ⇒ DOM 判据无从落地"

    def test_every_pending_entry_host_mounts_the_notice(self, manifest_slice: dict) -> None:
        """三要素：宿主 import + 模板挂载 + 传自己的 entry id，缺一即红。"""
        checked = 0
        for entry in manifest_slice["independent_entries"]:
            if entry.get("adapter_id") is not None:
                continue  # 已注册 adapter 的 entry 由真实 bridge 状态表达
            host = ROOT / entry["host_path"]
            assert host.is_file(), f"{entry['entry_id']} 的宿主不存在: {host}"
            source = _strip_ts_comments(host.read_text(encoding="utf-8"))
            assert f"import {NOTICE_COMPONENT_NAME} from" in source, (
                f"{host.name} 没有 import {NOTICE_COMPONENT_NAME}"
            )
            template = _vue_template(source)
            assert f"<{NOTICE_COMPONENT_NAME}" in template, (
                f"{host.name} 的模板里没有挂 <{NOTICE_COMPONENT_NAME}> ⇒ 结构性死代码"
            )
            mount_pattern = re.compile(
                r"<" + NOTICE_COMPONENT_NAME + r"\s+entry-id=\"([^\"]+)\"\s*/?>"
            )
            found = mount_pattern.findall(template)
            assert len(found) == 1, (
                f"{host.name} 的 {NOTICE_COMPONENT_NAME} 挂载点命中 {len(found)} 次（应为 1）"
            )
            assert found[0] == entry["entry_id"], (
                f"{host.name} 传的 entry-id={found[0]!r}，应为 {entry['entry_id']!r}"
            )
            checked += 1
        assert checked == 17, f"只校验到 {checked} 个宿主，应为 17 —— 分母缩水"

    def test_notice_mount_sits_inside_the_declared_mode_toolbar(
        self, manifest_slice: dict
    ) -> None:
        """通知必须挂在**该宿主自己声明的**模式工具栏区块内。

        🔴 工具栏门控锚点逐 entry 从 slice 的 `ui_toolbar_gate` 读：G 循环 17 个宿主的 class
        逐个不同，4 个宿主另有 `gN-index-toolbar` 兄弟区块。写死单一正则会让其余的区块判据
        恒真（挂载点不在被搜区块里也不报错）。
        """
        gates: set[str] = set()
        for entry in manifest_slice["independent_entries"]:
            gate = entry.get("ui_toolbar_gate")
            assert isinstance(gate, str) and gate.strip(), (
                f"{entry['entry_id']} 缺 ui_toolbar_gate ⇒ 区块判据无锚点"
            )
            gates.add(gate)
            host = ROOT / entry["host_path"]
            source = _strip_ts_comments(host.read_text(encoding="utf-8"))
            template = _vue_template(source)
            block = _toolbar_block(template, gate)
            assert f"<{NOTICE_COMPONENT_NAME}" in block, (
                f"{host.name} 的通知没挂在 {gate!r} 界定的模式工具栏区块内"
            )
            assert "el-segmented" in block, (
                f"{host.name} 的 {gate!r} 区块里没有模式切换器 ⇒ 它不是模式工具栏，锚点选错了"
            )
        assert len(gates) == 17, (
            f"17 个宿主的工具栏锚点应逐个不同，实测只有 {len(gates)} 个不同值 ⇒ "
            "有宿主的门控被误抄成别人的"
        )

    def test_ui_gate_source_refs_and_dom_guard_ref_are_real(
        self, manifest_slice: dict
    ) -> None:
        """范式 step 11 要求的 ui_gate_source_refs / dom_guard_ref 必须真存在。"""
        for entry in manifest_slice["independent_entries"]:
            refs = entry.get("ui_gate_source_refs")
            assert isinstance(refs, list) and len(refs) >= 3, (
                f"{entry['entry_id']} 的 ui_gate_source_refs 太少（宿主 + 组件 + 文案模块）"
            )
            for ref in refs:
                assert (ROOT / ref.split("#")[0]).exists(), (
                    f"{entry['entry_id']} 的 ui_gate_source_refs 指向不存在的路径 {ref!r}"
                )
            guard = entry.get("dom_guard_ref")
            assert isinstance(guard, str) and guard, f"{entry['entry_id']} 缺 dom_guard_ref"
            path, _, node = guard.partition("::")
            assert (ROOT / path).is_file(), f"dom_guard_ref 指向不存在的文件: {path}"
            body = (ROOT / path).read_text(encoding="utf-8")
            cls = node.split("::")[0]
            assert f"class {cls}" in body, (
                f"dom_guard_ref 指向的类 {cls} 不在 {path} 里 ⇒ 声明了但没人实现"
            )

    def test_registered_entry_ids_agree_with_the_slice(self, manifest_slice: dict) -> None:
        """双向锁：slice 的 adapter_id 与前端登记表必须互相印证。"""
        source = _strip_ts_comments(NOTICE_MODULE.read_text(encoding="utf-8"))
        block = re.search(
            r"SYNC_ADAPTER_REGISTERED_ENTRY_IDS:\s*readonly\s+string\[\]\s*=\s*\[([\s\S]*?)\]",
            source,
        )
        assert block, "找不到 SYNC_ADAPTER_REGISTERED_ENTRY_IDS 的声明"
        registered = set(re.findall(r"['\"]([^'\"]+)['\"]", block.group(1)))
        for entry in manifest_slice["independent_entries"]:
            if entry.get("adapter_id") is None:
                assert entry["entry_id"] not in registered, (
                    f"{entry['entry_id']} 在 slice 里没有 adapter，却被前端登记为已注册 ⇒ "
                    "界面会以「已双向」呈现"
                )
            else:
                assert entry["entry_id"] in registered, (
                    f"{entry['entry_id']} 在 slice 里已有 adapter，前端登记表却没跟上 ⇒ "
                    "界面会继续挂假警告"
                )

    def test_hosts_do_not_claim_bidirectional_writeback(self, manifest_slice: dict) -> None:
        """AC 1.4 前半句：宿主不得出现「可双向回写」「同步成功」之类的成功态宣称。

        注：G1 宿主有「拉取中…」「已拉取」「OO不可用」三个标签 —— 分别是进行态、
        config 预拉成功态与能力否认，都不是「双向回写成功」的宣称，不在禁列里。
        """
        forbidden = ("可双向回写", "双向同步", "同步成功", "已同步")
        for entry in manifest_slice["independent_entries"]:
            host = ROOT / entry["host_path"]
            source = _strip_ts_comments(host.read_text(encoding="utf-8"))
            for word in forbidden:
                assert word not in source, (
                    f"{host.name} 出现 {word!r} —— 未注册 adapter 的入口不得宣称双向/同步成功"
                )


# ════════════════════════════════════════════════════════════════════════════
# 范式一致性与源码结构
# ════════════════════════════════════════════════════════════════════════════
class TestParadigmCompliance:
    """验证 slice / deletion plan 与 Task 45 冻结的范式一致（范式 JSON 只读）。"""

    def test_paradigm_exists(self, paradigm: dict) -> None:
        assert paradigm["schema_version"] == "migration-paradigm:v1"
        assert paradigm["frozen_by"] == "Task 45"

    def test_task49_is_in_both_paradigm_scopes(self, paradigm: dict) -> None:
        """Task 49 同时受 definition_producer_paradigm 与 slice_schema 约束。"""
        producer = paradigm["definition_producer_paradigm"]
        assert 49 in producer["applies_to_tasks"]
        assert 49 in paradigm["slice_schema"]["applies_to_tasks"]
        step3 = next(s for s in producer["steps"] if s["step"] == 3)
        assert step3["name"] == "resolve_html_counterpart"
        assert set(step3["output"]["allowed_values"]) == set(HTML_COUNTERPART_VERDICTS)
        step11 = next(s for s in producer["steps"] if s["step"] == 11)
        assert step11["name"] == "enforce_honest_mode_visibility"
        assert set(step11["output"]["must_contain"]) == {"ui_gate_source_refs", "dom_guard_ref"}

    def test_capability_enum_comes_from_the_paradigm(self, paradigm: dict) -> None:
        """本文件的 CAPABILITY_ENUM 必须与范式登记的枚举一致（不另立一份真源）。"""
        assert set(paradigm["adjudication_criteria"]["capability_enum"]) == set(CAPABILITY_ENUM)

    def test_single_onlyoffice_illegal_criteria_are_registered(self, paradigm: dict) -> None:
        """范式明文把「无 adapter / 无 contract / 无 bundle」列为非法判据。"""
        verdict = paradigm["adjudication_criteria"]["verdicts"]["single_onlyoffice"]
        assert verdict["forbidden_verdict_value"] == "exists"
        assert verdict["illegal_criteria"], "illegal_criteria 为空 ⇒ AP-1 判据无分母"

    def test_slice_satisfies_the_schema_required_fields(
        self, manifest_slice: dict, paradigm: dict
    ) -> None:
        """逐项校验 slice_schema.required_* **加** effective_from_task_48 的追加必填。"""
        schema = paradigm["slice_schema"]
        assert schema["target_schema_version"] == manifest_slice["schema_version"]
        for key in schema["required_top_level"]:
            assert key in manifest_slice, f"slice 缺顶层字段 {key}"
        for key in schema["required_slice_scope"]:
            assert key in manifest_slice["slice_scope"], f"slice_scope 缺 {key}"
        for item in manifest_slice["slice_scope"]["excluded_from_slice"]:
            for key in schema["required_excluded_item"]:
                assert key in item
        for key in schema["required_authoritative_templates"]:
            assert key in manifest_slice["authoritative_templates"]
        for record in manifest_slice["authoritative_templates"]["files"]:
            for key in schema["required_template_file"]:
                assert key in record, f"模板 {record.get('name')} 缺 {key}"

        effective = schema["effective_from_task_48"]
        assert set(effective["required_entry"]) == {
            "html_counterpart_verdict",
            "html_counterpart_source_refs",
        }
        required_entry = list(schema["required_entry"]) + list(effective["required_entry"])
        for entry in manifest_slice["independent_entries"]:
            for key in required_entry:
                assert key in entry, f"{entry['entry_id']} 缺 entry 字段 {key}"
            for key in schema["presence_required_value_may_be_null"]:
                assert key in entry, f"{entry['entry_id']} 缺（值可为 null 的）字段 {key}"
            for key in schema["required_entry_adjudication"]:
                assert key in entry["adjudication"], (
                    f"{entry['entry_id']} 的 adjudication 缺 {key}"
                )
            for key in schema["required_entry_evidence"]:
                assert key in entry["evidence"], f"{entry['entry_id']} 的 evidence 缺 {key}"

        required_bp = list(schema["required_blocking_precondition"]) + list(
            effective["required_blocking_precondition"]
        )
        for bp in manifest_slice["blocking_preconditions"]:
            for key in required_bp:
                assert key in bp, f"阻断项 {bp.get('id')} 缺 {key}"
        for key in schema["required_cross_entry_isolation"]:
            assert key in manifest_slice["cross_entry_isolation"]
        for key in schema["required_honest_adjudication_summary"]:
            assert key in manifest_slice["honest_adjudication_summary"], (
                f"honest_adjudication_summary 缺 {key}"
            )
        for key in schema["required_slice_counters"]:
            assert key in manifest_slice["honest_adjudication_summary"]["slice_counters"]

    def test_dynamic_row_identity_section_satisfies_the_conditional_schema(
        self, manifest_slice: dict, paradigm: dict
    ) -> None:
        """conditional_sections 的 dynamic_row_identity 一旦出现就必须字段齐备。"""
        spec = next(
            cs
            for cs in paradigm["slice_schema"]["conditional_sections"]
            if cs["section"] == "dynamic_row_identity"
        )
        body = manifest_slice["dynamic_row_identity"]
        for key in spec["required_fields"]:
            assert key in body, f"dynamic_row_identity 缺 {key}"
        entry_ids = {e["entry_id"] for e in manifest_slice["independent_entries"]}
        assert len(body["tables"]) == len(entry_ids), (
            f"dynamic_row_identity 登记 {len(body['tables'])} 张表，entry 有 {len(entry_ids)} 条 "
            "⇒ 有 entry 的动态行表没登记"
        )
        for table in body["tables"]:
            for key in spec["required_table_fields"]:
                assert key in table, f"{table.get('table_key')} 缺 {key}"
            assert table["entry_id"] in entry_ids
            for key in spec["required_row_identity_fields"]:
                assert key in table["row_identity"], (
                    f"{table['table_key']}.row_identity 缺 {key}"
                )
            ref = table["row_identity"]["source_ref"]
            assert (ROOT / ref.split("#")[0]).is_file(), (
                f"{table['table_key']} 的 source_ref 指向不存在的文件 {ref!r}"
            )
            assert table["row_identity"]["kind"] not in body["forbidden_identity_kinds"], (
                f"{table['table_key']} 的 row_identity.kind 落在 forbidden_identity_kinds 里"
            )

    def test_the_slice_passes_the_paradigm_nominated_validator(
        self, manifest_slice: dict
    ) -> None:
        """被范式点名的校验器对本 slice **一条违规都不许报**。

        🔴 判据换过一次，换的理由要留痕：本条原名
        `test_the_only_schema_violations_are_the_registered_sr3_gap`，断言「违规存在且**全部**
        是 SR-3」—— 那是缺口 G1（同一条 SR-3 有两个互相矛盾的判官）还没收口时的诚实形态，
        它自己的 docstring 写明「若范式 owner 收口了 SR-3 ⇒ 违规变空 ⇒ 打红，提醒回来删登记」。
        范式 owner 已把 SR-3 改成蕴含式（`capability` 在 `capability_enum` 内 **或** 为 null 且
        `required_pending_verdict_fields` 三字段齐备）并新增 SR-9（待裁决条目数 ==
        `slice_counters.unadjudicated`），待裁决态成为合法一等公民 ⇒ 本 slice 的 17 条
        `capability: null` 不再是违规，旧判据按其设计当场打红。按新事实重写为「通过」。

        两个方向仍然都咬：
        * 本 slice 破了任何 schema 规则（含把待裁决态写得不齐备：缺 `capability_verdict_stage`
          / `capability_target` 不在枚举内 / `capability_target_blocked_by` 空数组，或
          `unadjudicated` 计数说谎）⇒ 打红；
        * 校验器被短路成恒返回空列表 ⇒ 由下面两条对合成残缺 slice 的反向自检打红。

        为什么当初不能靠「给 capability 填一个枚举值」消掉那条红：违反 AC 12.8（有 HTML 对端
        不得 `single_onlyoffice`）或 AC 12.1 + SR-6（五个身份字段全 null 不得 `bidirectional`）。
        收口走的是范式侧，不是往 slice 里填假裁决。
        """
        contract = pathlib.Path(__file__).with_name("test_migration_paradigm_contract.py")
        assert contract.is_file(), f"{contract} 不存在 —— 被点名的校验器缺失"
        import copy  # noqa: PLC0415
        import importlib.util  # noqa: PLC0415

        spec = importlib.util.spec_from_file_location("_t49_paradigm_contract_host", contract)
        assert spec and spec.loader
        module = importlib.util.module_from_spec(spec)
        spec.loader.exec_module(module)

        problems = module.validate_slice_against_schema(manifest_slice)
        assert problems == [], (
            f"本 slice 通不过范式自己的 slice_schema（{len(problems)} 条）：\n"
            + "\n".join(f"  !! {p}" for p in problems)
        )
        # 反向自检 ①：SR-3 右支不是「null 一律通过」。抽掉第一条 entry 的待裁决声明一项，
        # 校验器必须当场点名该字段。
        crippled = copy.deepcopy(manifest_slice)
        crippled["independent_entries"][0].pop("capability_verdict_stage")
        crippled_problems = module.validate_slice_against_schema(crippled)
        assert any("capability_verdict_stage" in p for p in crippled_problems), (
            "抽掉 capability_verdict_stage 后校验器不点名它 ⇒ 待裁决态被放宽成「null 一律通过」，"
            f"本条的「通过」是假绿。实际违规：{crippled_problems}"
        )
        # 反向自检 ②：SR-9 真在管计数。把 unadjudicated 改成说谎值，必须报违规。
        declared_pending = manifest_slice["honest_adjudication_summary"]["slice_counters"][
            "unadjudicated"
        ]
        lying = copy.deepcopy(manifest_slice)
        lying["honest_adjudication_summary"]["slice_counters"]["unadjudicated"] = (
            declared_pending + 1
        )
        lying_problems = module.validate_slice_against_schema(lying)
        assert lying_problems, (
            "把 slice_counters.unadjudicated 改成说谎值后校验器一条都不报 ⇒ SR-9 没在管计数，"
            "SR-3 右支就成了后门（把全部 entry 记成待裁决同时报 unadjudicated=0 也能过）"
        )
        # 正向：待裁决条目数与计数现算相等（SR-9 的本 slice 实例）
        pending = [e for e in manifest_slice["independent_entries"] if e.get("capability") is None]
        assert len(pending) == declared_pending == len(manifest_slice["independent_entries"]), (
            f"SR-9 现算不等：待裁决 entry {len(pending)} 条 / slice_counters.unadjudicated "
            f"{declared_pending} / entry 总数 {len(manifest_slice['independent_entries'])}"
        )
        # 收口留痕：slice 自述里的 G1 登记必须已改写成「已解除」，不许留着过期的「已知红」。
        registered = manifest_slice["paradigm_schema_conflict"]["known_red_it_causes"]
        assert registered["status"] == "resolved", (
            "G1 已收口但 slice 的 known_red_it_causes.status 不是 resolved ⇒ 自述过期"
        )
        assert "SR-3" in registered["resolved_by"] and "SR-9" in registered["resolved_by"], (
            "known_red_it_causes.resolved_by 必须写明是哪两条规则接纳并管住了待裁决态"
        )
        assert not registered["failing_test"], (
            "G1 已解除，known_red_it_causes.failing_test 必须清空 —— 留着旧 nodeid 会让下一个人"
            "以为那条测试还在红"
        )
        assert "test_the_slice_passes_the_paradigm_nominated_validator" in (
            registered["locked_by"]
        ), "known_red_it_causes.locked_by 必须指向重写后的判据名"

    def test_paradigm_conflict_is_registered(self, manifest_slice: dict) -> None:
        """范式内部冲突必须登记，不许悄悄绕过。"""
        conflict = manifest_slice["paradigm_schema_conflict"]
        for key in (
            "what",
            "how_resolved_here",
            "residual_inconsistency",
            "sibling_slice_conflicts",
        ):
            assert len(conflict.get(key, "")) > 40, f"paradigm_schema_conflict 缺 {key}"

    def test_deletion_plan_has_paradigm_ref(self, deletion_plan: dict) -> None:
        assert deletion_plan["paradigm_ref"] == (
            "backend/data/workpaper_sync_migration_paradigm.json"
        )
        assert deletion_plan["adjudication_source_of_truth"].startswith(
            "backend/data/workpaper_sync_g_cycle_manifest_slice.json"
        )
        assert deletion_plan["deletion_execution_owner"]

    def test_properties_and_requirements_are_declared(
        self, manifest_slice: dict, deletion_plan: dict
    ) -> None:
        """slice 与 deletion plan 声明的 Property / Requirement 必须与任务正文一致。"""
        expected_props = {20, 21, 28, 69, 70}
        assert set(manifest_slice["properties_verified"]) == expected_props
        assert set(deletion_plan["properties_verified"]) == expected_props
        expected_reqs = {
            "6.1", "6.2", "6.10", "9.1", "12.1", "12.4",
            "12.10", "12.11", "12.12", "14.1",
        }
        assert expected_reqs <= set(manifest_slice["requirements_covered"]), (
            f"slice 的 requirements_covered 缺 "
            f"{sorted(expected_reqs - set(manifest_slice['requirements_covered']))}"
        )


class TestSourceCodeStructure:
    """验证源码结构与 slice/deletion plan 一致（现状锁，删除归 Task 66/72）。"""

    def test_hosts_exist_and_still_import_their_legacy_composable(
        self, manifest_slice: dict
    ) -> None:
        """宿主存在且仍引用各自的 legacy dual-mode composable（BP-12 登记的未删状态）。"""
        seen: set[str] = set()
        for entry in manifest_slice["independent_entries"]:
            host = ROOT / entry["host_path"]
            assert host.is_file()
            source = _strip_ts_comments(host.read_text(encoding="utf-8"))
            expected = entry["legacy_dual_mode_composable"]
            assert re.search(
                r"import\s*\{?\s*" + expected + r"\b", source
            ), f"{host.name} 应仍 import {expected}"
            assert expected not in seen, (
                f"{expected} 被两个宿主共用 ⇒ 与 slice 的逐 entry 独立结论矛盾"
                "（唯一合法例外是 useG1DualMode 被 E 循环共用，那是跨 slice 不是本 slice 内）"
            )
            seen.add(expected)
        assert len(seen) == 17

    def test_g_cycle_templates_exist(self) -> None:
        g_dir = TEMPLATE_DIR / "G"
        assert g_dir.is_dir()
        xlsx_files = [p for p in g_dir.glob("G*.xlsx") if not p.name.startswith("~$")]
        assert len(xlsx_files) == 15, f"G 循环应有 15 个 xlsx 模板，实际 {len(xlsx_files)}"

    def test_bridge_adapter_supports_entry_id(self) -> None:
        adapter = SYNC_DIR / "usePilotBridgeAdapter.ts"
        assert adapter.is_file(), "usePilotBridgeAdapter.ts 必须存在（Task 45 产物）"
        assert "entryId" in adapter.read_text(encoding="utf-8")

# -*- coding: utf-8 -*-
"""D4 位置数组（positional array）往返等值 —— Wave 1 阻塞门。

spec: d4-revenue-matrix-bidirectional · Task 2 / 2.1
Requirements: 1.1–1.6 · Properties 1 / 2 / 6

必须走真实 ``parse_contract → build_excel_adapter → materialize → extract →
merge_projection_into_store_rows``；provider 私有 helper 不得单独解除本门。
"""
from __future__ import annotations

import hashlib
import io
import json
import os
import sys
import uuid
from pathlib import Path
from typing import Any, Mapping

import pytest
from hypothesis import given, settings, strategies as st
from openpyxl import Workbook

_REPO = Path(__file__).resolve().parents[3]
_BACKEND = _REPO / "backend"
if str(_BACKEND) not in sys.path:  # pragma: no cover
    sys.path.insert(0, str(_BACKEND))
os.environ.setdefault("DB_DISABLE_SSL", "True")

from app.services.excel_structure_fingerprint import identity_inventory  # noqa: E402
from app.services.workpaper_sync import excel_extract as X  # noqa: E402
from app.services.workpaper_sync import excel_instrumentation as EI  # noqa: E402
from app.services.workpaper_sync.adapters import excel as AX  # noqa: E402
from app.services.workpaper_sync.adapters.base import FieldValue, Projection  # noqa: E402
from app.services.workpaper_sync.contracts import (  # noqa: E402
    CONTRACT_SCHEMA_VERSION,
    parse_contract,
)
from app.services.workpaper_sync.excel_entry_gate import (  # noqa: E402
    AdapterBuild,
    FrozenEntryDefinitions,
    parse_identity_inventory,
)
from app.services.workpaper_sync.json_path import (  # noqa: E402
    JsonPathArrayIndexOobError,
    JsonPathArrayLengthInvalidError,
    JsonPathMissingSegmentError,
    JsonPathTypeMismatchError,
    resolve_json_path,
    set_json_path,
)
from app.services.workpaper_sync.models import (  # noqa: E402
    AuthorityModel,
    BundleSlot,
    BundleSlotSpec,
    DefinitionState,
)
from app.services.workpaper_sync.resolution import DefinitionBundleSnapshot  # noqa: E402

ENTRY_ID = "d4.revenue_detail.fixture"
CONTRACT_ID = "d4.revenue_detail.fixture"
ADAPTER_ID = "d4.revenue_detail.fixture"
MANAGED_SHEET = "主营业务收入明细表D4-2"
TABLE_NAME = "GT_D42_ROWS"
ROWS_TABLE_KEY = "revenue_detail_rows"
UUID_COL = "N"
FIRST_ROW = 2
LAST_ROW = 2
FOOTER_ROW = 3
MONTH_COLS = tuple("BCDEFGHIJKLM")
ROW_IDENTITY_STORE_KEY = "rowId"
TEMPLATE_SHA = hashlib.sha256(b"d4-positional-array-fixture-v1").hexdigest()


def _d(label: str) -> str:
    return hashlib.sha256(label.encode("utf-8")).hexdigest()


def _stable_key(column_key: str, row_identity: str = "{row_uuid}") -> str:
    return f"{ROWS_TABLE_KEY}/{row_identity}/{column_key}"


def _month_field_specs() -> list[tuple[str, str, str]]:
    return [
        (f"month_{i + 1:02d}", col, f"months/{i}")
        for i, col in enumerate(MONTH_COLS)
    ]


def _minimal_workbook_bytes() -> bytes:
    wb = Workbook()
    ws = wb.active
    ws.title = MANAGED_SHEET
    ws["A1"] = "项目"
    for i, col in enumerate(MONTH_COLS):
        ws[f"{col}1"] = f"{i + 1}月"
    ws["A2"] = "产品A"
    for col in MONTH_COLS:
        ws[f"{col}2"] = 0
    ws["A3"] = "合计"
    buf = io.BytesIO()
    wb.save(buf)
    return buf.getvalue()


SPEC = EI.ExcelInstrumentationSpec(
    entry_id=ENTRY_ID,
    template_id="D42",
    template_relative_path="D/D4 收入底稿.xlsx",
    managed_sheet=MANAGED_SHEET,
    first_data_row=FIRST_ROW,
    last_data_row=LAST_ROW,
    footer_row=FOOTER_ROW,
    managed_last_col="M",
    uuid_col=UUID_COL,
    table_name=TABLE_NAME,
)


def contract_payload() -> dict[str, Any]:
    fields = [
        {
            "stable_field_key": _stable_key(column_key),
            "json_pointer": f"/rows/{{row_uuid}}/{json_path}",
            "column_key": column_key,
            "cell": {"column": col, "row_from": "row_identity"},
            "mode": "editable",
            "value_type": "amount",
            "source_ref": f"源xlsx!{MANAGED_SHEET}!{col}{FIRST_ROW}",
        }
        for column_key, col, json_path in _month_field_specs()
    ]
    rows_table: dict[str, Any] = {
        "table_key": ROWS_TABLE_KEY,
        "anchor": f"A{FIRST_ROW}",
        "header_rows": 1,
        "row_identity": {
            "kind": "field",
            "json_pointer": f"/rows/*/{ROW_IDENTITY_STORE_KEY}",
        },
        "delete_policy": "tombstone",
        "footer_anchor": {"marker": "合计", "search_column": "A"},
        "fields": fields,
    }
    return {
        "schema_version": CONTRACT_SCHEMA_VERSION,
        "contract_id": CONTRACT_ID,
        "semantic_version": "1.0.0",
        "review_status": "reviewed",
        "document_type": "xlsx",
        "template_definition_sha256": _d("d4-fixture-template-definition"),
        "instrumentation_definition_sha256": _d("d4-fixture-instrumentation-definition"),
        "template": {
            "relative_path": "D/D4 收入底稿.xlsx",
            "template_sha256": TEMPLATE_SHA,
            "normalized_structure_hash": _d("d4-fixture-normalized-structure"),
        },
        "identity_carriers": [
            "hidden_sheet",
            "defined_name",
            "excel_table",
            "hidden_uuid_column",
        ],
        "sheets": [
            {
                "sheet_key": "d42-managed",
                "excel_name": MANAGED_SHEET,
                "locator": {"anchor": X.TABLE_SHEET_ANCHOR},
                "tables": [rows_table],
            }
        ],
    }


def make_bundle(contract: Any) -> DefinitionBundleSnapshot:
    def slot(kind: BundleSlot, digest: str) -> BundleSlotSpec:
        return BundleSlotSpec(kind, "definition", f"definition:{uuid.uuid4()}", digest)

    return DefinitionBundleSnapshot(
        bundle_id=uuid.uuid4(),
        bundle_sha256=_d("d4-fixture-bundle"),
        schema_version="definition-bundle:v1",
        state=DefinitionState.approved,
        authority_model=AuthorityModel.projection_contract,
        authority_model_definition_id=uuid.uuid4(),
        authority_model_definition_sha256=_d("d4-fixture-authority"),
        slots={
            BundleSlot.template: slot(
                BundleSlot.template, contract.template_definition_sha256
            ),
            BundleSlot.instrumentation: slot(
                BundleSlot.instrumentation, contract.instrumentation_definition_sha256
            ),
            BundleSlot.contract: slot(BundleSlot.contract, contract.canonical_sha256),
        },
    )


def make_definitions(
    contract: Any, inventory_raw: Mapping[str, Any]
) -> FrozenEntryDefinitions:
    return FrozenEntryDefinitions(
        entry_id=ENTRY_ID,
        bundle=make_bundle(contract),
        contract=contract,
        adapter_build=AdapterBuild(
            adapter_id=ADAPTER_ID,
            adapter_build_digest=_d("d4-fixture-adapter-build"),
            document_type="xlsx",
            contract_version="1.0.0",
        ),
        identity_inventory=parse_identity_inventory(inventory_raw),
        business_sheets=(),
        dynamic_column_keys={},
        structure_inventory_size=0,
    )


BINDING = X.ExcelIdentityBinding(
    table_name=TABLE_NAME, uuid_column=UUID_COL, table_key=ROWS_TABLE_KEY
)


def merge_projection_into_store_rows(
    *,
    projection: Any,
    base_rows: list[Mapping[str, Any]],
) -> tuple[list[dict[str, Any]], int, int, set[str]]:
    """fixture 本地 merge —— 强制走共享 ``set_json_path``（数组感知）。"""
    field_to_path = {ck: jp for ck, _col, jp in _month_field_specs()}
    by_id: dict[str, dict[str, Any]] = {}
    order: list[str] = []
    for row in base_rows:
        rid = str(row.get(ROW_IDENTITY_STORE_KEY) or "").strip()
        if not rid:
            continue
        by_id[rid] = dict(row)
        months = by_id[rid].get("months")
        if isinstance(months, list):
            by_id[rid]["months"] = list(months)
        order.append(rid)

    applied = 0
    visited = 0
    touched: set[str] = set()
    for key in projection.stable_keys():
        fv = projection.get(key)
        if fv is None:
            continue
        rid = getattr(fv, "row_key", None)
        if not rid:
            continue
        target = by_id.setdefault(str(rid), {ROW_IDENTITY_STORE_KEY: str(rid)})
        if str(rid) not in order:
            order.append(str(rid))
        field_id = str(key).rsplit("/", 1)[-1]
        json_path = field_to_path.get(field_id)
        if not json_path:
            continue
        visited += 1
        if set_json_path(target, json_path, getattr(fv, "value", None)):
            applied += 1
            touched.add(str(rid))
    return [by_id[rid] for rid in order], applied, visited, touched


def _projection_from_store_row(
    row: Mapping[str, Any], *, contract: Any, row_identity: str
) -> Projection:
    values: dict[str, FieldValue] = {}
    for column_key, _col, json_path in _month_field_specs():
        spec = contract.field_by_stable_key(_stable_key(column_key))
        value = resolve_json_path(row, json_path)
        sk = _stable_key(column_key, row_identity)
        values[sk] = FieldValue(
            stable_key=sk,
            value=value,
            value_type=spec.value_type,
            mode=spec.mode,
            row_key=row_identity,
        )
    return Projection(
        contract_id=contract.contract_id,
        semantic_version=contract.semantic_version,
        document_type=contract.document_type,
        values=values,
        row_keys={ROWS_TABLE_KEY: (row_identity,)},
    )


@pytest.fixture(scope="module")
def instrumented() -> EI.InstrumentedWorkbook:
    gate = EI.ExcelIdentityCarrierGate.load()
    return EI.instrument_workbook_bytes(_minimal_workbook_bytes(), SPEC, gate=gate)


@pytest.fixture(scope="module")
def contract() -> Any:
    return parse_contract(contract_payload(), adapter_id=ADAPTER_ID)


@pytest.fixture(scope="module")
def definitions(contract: Any, instrumented: EI.InstrumentedWorkbook):
    inv = identity_inventory(
        instrumented.instrumented_bytes,
        expected_table=TABLE_NAME,
        uuid_column_letter=UUID_COL,
    )
    return make_definitions(contract, inv)


@pytest.fixture
def base_path(instrumented: EI.InstrumentedWorkbook, tmp_path: Path) -> Path:
    path = tmp_path / "base.xlsx"
    path.write_bytes(instrumented.instrumented_bytes)
    return path


class TestParseContractAcceptsArrayIndexPointer:
    """Requirement 1.1"""

    def test_months_index_pointer_accepted(self, contract: Any) -> None:
        for i in range(12):
            pointer = f"/rows/{{row_uuid}}/months/{i}"
            field = contract.field_by_stable_key(_stable_key(f"month_{i + 1:02d}"))
            assert field.json_pointer == pointer
            assert field.json_pointer.rsplit("/", 1)[-1] == str(i)


class TestSharedJsonPathFailClosed:
    """Requirements 1.3 / 1.4 · Property 2"""

    def test_oob_months_12_raises(self) -> None:
        row = {"months": [0] * 12}
        with pytest.raises(JsonPathArrayIndexOobError) as exc:
            resolve_json_path(row, "months/12")
        assert exc.value.error_code == "json_path_array_index_oob"
        with pytest.raises(JsonPathArrayIndexOobError) as exc2:
            set_json_path(row, "months/12", 1)
        assert exc2.value.error_code == "json_path_array_index_oob"

    def test_length_not_12_raises(self) -> None:
        row = {"months": [0] * 11}
        with pytest.raises(JsonPathArrayLengthInvalidError) as exc:
            resolve_json_path(row, "months/0")
        assert exc.value.error_code == "json_path_array_length_invalid"
        with pytest.raises(JsonPathArrayLengthInvalidError) as exc2:
            set_json_path({"months": [0] * 13}, "months/0", 9)
        assert exc2.value.error_code == "json_path_array_length_invalid"

    def test_missing_and_type_mismatch(self) -> None:
        with pytest.raises(JsonPathMissingSegmentError) as exc:
            resolve_json_path({}, "months/0")
        assert exc.value.error_code == "json_path_missing_segment"
        with pytest.raises(JsonPathTypeMismatchError) as exc2:
            resolve_json_path({"months": {"0": 1}}, "months/0")
        assert exc2.value.error_code == "json_path_type_mismatch"
        row = {"months": [0] * 12}
        set_json_path(row, "months/3", 7)
        assert isinstance(row["months"], list)
        assert row["months"][3] == 7


class TestPositionalArrayProductionRoundtrip:
    """Requirements 1.2 / 1.5 / 1.6 · Properties 1 / 6"""

    def test_materialize_extract_merge_keeps_months_list(
        self,
        contract: Any,
        definitions: FrozenEntryDefinitions,
        base_path: Path,
        instrumented: EI.InstrumentedWorkbook,
        tmp_path: Path,
    ) -> None:
        row_uuid = instrumented.row_uuids[FIRST_ROW]
        months = [float(i * 10 + 1) for i in range(12)]
        store_row = {
            ROW_IDENTITY_STORE_KEY: row_uuid,
            "product": "产品A",
            "months": list(months),
        }
        projection = _projection_from_store_row(
            store_row, contract=contract, row_identity=row_uuid
        )

        adapter = AX.build_excel_adapter(
            definitions=definitions, binding=BINDING, direction="html_to_oo"
        )
        staged = tmp_path / AX.STAGING_NAMESPACE / "staged.xlsx"
        staged.parent.mkdir(parents=True, exist_ok=True)
        result = adapter.materialize(
            substrate=base_path,
            projection=projection,
            output=staged,
            contract=definitions.contract,
        )
        assert staged.is_file() and staged.read_bytes()[:2] == b"PK"
        assert result.managed_field_count == 12

        extracted = adapter.extract(artifact=staged, contract=definitions.contract)
        for i in range(12):
            sk = _stable_key(f"month_{i + 1:02d}", row_uuid)
            fv = extracted.get(sk)
            assert fv is not None, f"extract 缺失 {sk}"
            assert float(fv.value) == months[i]

        base_for_merge = [
            {
                ROW_IDENTITY_STORE_KEY: row_uuid,
                "product": "产品A",
                "months": [0.0] * 12,
            }
        ]
        merged, applied, visited, touched = merge_projection_into_store_rows(
            projection=extracted, base_rows=base_for_merge
        )
        assert visited == 12 and applied == 12
        assert touched == {row_uuid}
        assert len(merged) == 1
        out = merged[0]
        assert isinstance(out["months"], list)
        assert len(out["months"]) == 12
        assert [float(x) for x in out["months"]] == months
        serialized = json.dumps(out["months"], ensure_ascii=False)
        assert serialized.startswith("[")
        assert not serialized.lstrip().startswith("{")


class TestPropertyMinimalWriteSurface:
    """Task 2.1 PBT · Requirement 1.2"""

    @given(
        months=st.lists(
            st.floats(
                min_value=-1e6, max_value=1e6, allow_nan=False, allow_infinity=False
            ),
            min_size=12,
            max_size=12,
        ),
        index=st.integers(min_value=0, max_value=11),
        new_value=st.floats(
            min_value=-1e6, max_value=1e6, allow_nan=False, allow_infinity=False
        ),
    )
    @settings(max_examples=40, deadline=None)
    def test_only_target_index_mutated(
        self, months: list[float], index: int, new_value: float
    ) -> None:
        row = {"months": list(months)}
        before = list(row["months"])
        changed = set_json_path(row, f"months/{index}", new_value)
        assert isinstance(row["months"], list)
        assert len(row["months"]) == 12
        for i, val in enumerate(row["months"]):
            if i == index:
                assert val == new_value
            else:
                assert val == before[i]
        assert changed == (before[index] != new_value)

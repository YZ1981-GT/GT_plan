# -*- coding: utf-8 -*-
"""N1 批 C 的 store item、载荷列与模板几何语义守卫。"""
from __future__ import annotations

import json
from pathlib import Path

import openpyxl

from app.services.workpaper_sync import phase5_n1_deferred_tax_assets as provider
from app.services.workpaper_sync.store_item_registry import STORE_MERGE_REGISTRY
from app.services.workpaper_sync.store_payload_columns import CONCLUSION_ONLY, payload_columns_of

_BACKEND = Path(__file__).resolve().parents[2]
_TEMPLATE = _BACKEND / "wp_templates" / provider.TEMPLATE_RELATIVE_PATH
_FRONTEND_DETAIL = (
    _BACKEND.parent
    / "audit-platform/frontend/src/components/workpaper/composables/useN1Detail.ts"
)


def test_n1_store_item_matches_frontend_contract_and_registry() -> None:
    expected = "N1-2-detail-rows"
    plan = STORE_MERGE_REGISTRY[provider.ADAPTER_ID]
    payload = provider.build_contract_payload()
    table = payload["sheets"][0]["tables"][0]
    assert provider.STORE_ITEM_ID == expected
    assert provider.all_store_item_ids() == (expected,)
    assert provider.SPEC_N12.store_item_id == expected
    assert plan.item_ids == (expected,)
    assert payload["review"]["html_store"]["item_id"] == expected
    assert {field["store_item_id"] for field in table["fields"]} == {expected}
    assert "const ITEM_PREFIX = 'N1-2-detail'" in _FRONTEND_DETAIL.read_text(encoding="utf-8")

    row = {
        "id": "n1-row",
        "itemName": "可抵扣亏损",
        "category": "可抵扣亏损",
        "beginDiff": 100,
        "beginTaxRate": 0.25,
        "beginAje": 1,
        "beginRje": 2,
        "endDiff": 120,
        "endTaxRate": 0.25,
        "endAje": 3,
        "endRje": 4,
    }
    projection = provider.build_store_projection(
        json.dumps([row], ensure_ascii=False),
        contract=provider.load_contract_from_disk(),
    )
    assert projection.row_keys[provider.ROWS_TABLE_KEY] == ("n1-row",)


def test_n1_payload_column_is_conclusion_at_runtime() -> None:
    assert provider.PAYLOAD_COLUMN == "conclusion"
    assert provider.PAYLOAD_COLUMN_MODE == "conclusion_only"
    assert payload_columns_of(provider) is CONCLUSION_ONLY


def test_n1_geometry_matches_authoritative_template() -> None:
    wb = openpyxl.load_workbook(_TEMPLATE, read_only=True, data_only=False)
    try:
        ws = wb[provider.MANAGED_SHEET]
        assert provider.FIRST_DATA_ROW == 11
        assert provider.LAST_DATA_ROW == 40
        assert provider.FOOTER_ROW == 42
        assert provider.SPEC_N12.footer_carries_total_formula is True
        assert ws["A40"].value == "……"
        assert ws["E40"].value == "=ROUND(C40*D40,2)"
        assert ws["N40"].value == "=K40+L40+M40"
        assert ws["C42"].value == "=SUM(C11:C41)"
        assert ws["N42"].value == "=SUM(N11:N41)"
        assert provider.FORMULA_MASK == (
            "E11:E40",
            "H11:H40",
            "K11:K40",
            "N11:N40",
        )
    finally:
        wb.close()

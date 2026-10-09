# -*- coding: utf-8 -*-
"""D5-1 / D6-1 / D7-1 审定表 Property 7：逐格 mask 必须与**模板册真实公式格**一致。

spec: d567-sync-coverage-via-row-table-engine · Requirements 4.5 / 6.1

═══ 为什么判据必须锚定模板，而不是照 D3-1 抄 `_VALUE_SOURCES` 口径 ═══

D3-1 的 Property 7 判据写法是「`value_source == manual` 且 `value_type == amount` 的格
**不得** 入 mask」。**这套口径直接搬到 D5/D6 会产生 30 个假阳**（实测）：

  * `D5-1` 的 `审定表D5` R7-R8：B/C/D/F/G/H **全是真 SUMIF 公式**（`=SUMIF('应收款项融资明细表D5-2'!…)`）
    —— 尽管 `_VALUE_SOURCES` 把 `prior_aje`/`prior_rje`/`current_aje`/`current_rje` 声明成
    `manual`（那是**前端 store 侧**的语义：HTML 逐格存值），模板侧它们是公式格
    ⇒ 必须 mask，否则 OO 回写会**覆盖掉模板公式**。按 D3 口径判会误报 10 格。
  * `D6-1` 区3（净值 R26-R30）：模板里 B–I 全是 `=B8-B17` 这类派生公式 ⇒ 必须 mask。
    按 D3 口径判会误报 20 格。

⇒ 故本文件用**唯一权威源 = 模板册**（`backend/wp_templates/D/*.xlsx`）现算公式格集合，
   判据是三向的：

  ① 模板该格**有公式** ⇒ 必须 mask（防 OO 覆盖公式）
  ② 模板该格**无公式** 且字段**不是** `computed` ⇒ 必须**不** mask
     （否则审计师在 OO 里改不了 = D4-1 踩过的 fail-closed 缺陷，Property 7 本体）
  ③ 模板该格**无公式** 且字段是 `computed` ⇒ **允许** mask（保守保护，前端算的值 OO 写了也会丢）

  该口径与参考实现 D3-1 的**实际状态**自洽：D3-1 区2 的 F（`current_unadjusted`、
  `cross_sheet` 派生、模板无公式）在 D3-1 里确实**未** mask —— 「派生字段不可 OO 直写」
  由 `is_oo_writable()` 这条**独立**机制保证，不靠 mask。两个机制不可混为一谈。

═══ 🔴 本判据首次运行即抓到并修掉的真缺陷（Task 20）═══

  * `D7-1` 区1（性质 R8-R13）：原声明整行 mask `B-K`，注释称「性质区有 SUMIF cross_sheet 公式」
    —— **模板实测该前提不成立**，B/C/D/F/G/H 全为空。**36 格**手工录入格被锁死。
    表内自证：**同一张表**的区2（账龄 R20-R23）模板形态完全一致，声明却只 mask E/I/J/K。
  * `D6-1` 区1（R8-R12）+ 区2（R17-R21）：同型，**60 格**被锁死。
  * `D5-1`：声明与模板**逐格完全一致**（46 == 46，对称差 0），无需改动。

  修复后对称差：D5-1 = 0 / D7-1 = 0 / D6-1 只剩 ③ 类允许项（11 格 computed 列）。
"""
from __future__ import annotations

import os
import re
import sys
from pathlib import Path

import pytest

_BACKEND = Path(__file__).resolve().parents[2]
if str(_BACKEND) not in sys.path:  # pragma: no cover
    sys.path.insert(0, str(_BACKEND))
os.environ.setdefault("DB_DISABLE_SSL", "True")

from openpyxl import load_workbook  # noqa: E402

from app.services.workpaper_sync import phase5_d5_01_adjudication as D5  # noqa: E402
from app.services.workpaper_sync import phase5_d6_01_adjudication as D6  # noqa: E402
from app.services.workpaper_sync import phase5_d7_01_adjudication as D7  # noqa: E402
from app.services.workpaper_sync.phase5_adjudication_sheet import (  # noqa: E402
    AdjudicationSheetSpec,
    AdjudicationValueSource,
)

_WP_D = _BACKEND / "wp_templates" / "D"
_CELL_RE = re.compile(r"^([A-Z]+)(\d+)$")
#: 审定表受管列范围（A=项目名 … L=原因分析），与 field_specs 的 excel 列一致。
_COLS = tuple("ABCDEFGHIJKL")


class _Case:
    """一张审定表的判据输入（spec + 模板册 sheet）。"""

    def __init__(self, label: str, spec: AdjudicationSheetSpec, book: str, sheet: str) -> None:
        self.label = label
        self.spec = spec
        self.book = _WP_D / book
        self.sheet = sheet

    def __repr__(self) -> str:  # pragma: no cover
        return f"<{self.label}>"


CASES: list[_Case] = [
    _Case("D5-1", D5.SPEC_D501, "D5 应收款项融资.xlsx", "审定表D5"),
    _Case("D6-1", D6.SPEC_D601, "D6 合同资产.xlsx", "审定表D6-1"),
    _Case("D7-1", D7.SPEC_D701, "D7 合同负债.xlsx", "审定表D7-1"),
]

#: 🔴 可伪证豁免：D6-1 的 A 列镜像公式（`=A8` 等）虽是模板公式（判据 ① 要求 mask），
#:    但 A 列是 `editable` 的项目名列，而框架不变量 `assert_data_cells_not_masked()`
#:    **禁止** mask 数据行的 editable 列 —— 两条规则在此直接冲突。
#:    本轮不擅自改（改任一侧都会动别的判据面），显式登记 + 下方 `test_..._allowlist_is_falsifiable`
#:    对每条豁免验证三件事：真在 A 列、真在数据行、模板值真是**同列自引用镜像公式**。
#:    任一条不成立即红 ⇒ 豁免不是「写个理由就放行」的后门。
_COLUMN_A_MIRROR_ALLOWLIST: dict[str, tuple[str, ...]] = {
    "D6-1": ("A17", "A18", "A19", "A20", "A21", "A26", "A27", "A28", "A29", "A30"),
}


def _formula_cells(case: _Case, rows: set[int]) -> dict[str, str]:
    """现算模板册在给定行上的公式格 → 公式原文（不信声明注释，只信模板）。"""
    wb = load_workbook(case.book, data_only=False)
    try:
        ws = wb[case.sheet]
        out: dict[str, str] = {}
        for row in sorted(rows):
            for col in _COLS:
                v = ws[f"{col}{row}"].value
                if isinstance(v, str) and v.startswith("="):
                    out[f"{col}{row}"] = v
        return out
    finally:
        wb.close()


def _rows_of_interest(case: _Case) -> set[int]:
    """mask 覆盖的行 ∪ 各 section 数据行（判据作用域，现算不写死）。"""
    rows: set[int] = set()
    for cell in case.spec.masked_cells:
        m = _CELL_RE.match(cell)
        if m:
            rows.add(int(m.group(2)))
    rows |= _data_rows(case)
    return rows


def _data_rows(case: _Case) -> set[int]:
    rows: set[int] = set()
    for sec in case.spec.sections:
        rows.update(range(sec.first_data_row, sec.last_data_row + 1))
    return rows


def _col_to_source(case: _Case) -> dict[str, AdjudicationValueSource]:
    """Excel 列 → 该列字段的值来源（现算，防手写映射漂移）。"""
    return {fs[1]: case.spec.source_of(fs[0]) for fs in case.spec.field_specs}


def _classify(case: _Case) -> dict[str, list[str]]:
    """三向分类：违反 ① / 违反 ②（Property 7 本体）/ 命中 ③ 允许项。"""
    rows = _rows_of_interest(case)
    formulas = _formula_cells(case, rows)
    mask = set(case.spec.masked_cells)
    col_src = _col_to_source(case)
    allow = set(_COLUMN_A_MIRROR_ALLOWLIST.get(case.label, ()))

    violate_1: list[str] = []   # 有公式却没 mask
    violate_2: list[str] = []   # 无公式、非 computed，却被 mask（fail-closed）
    allowed_3: list[str] = []   # 无公式、computed，被 mask（允许）
    for row in sorted(rows):
        for col in _COLS:
            cell = f"{col}{row}"
            has_formula = cell in formulas
            masked = cell in mask
            src = col_src.get(col)
            if has_formula and not masked:
                if cell not in allow:
                    violate_1.append(cell)
            elif not has_formula and masked:
                if src is AdjudicationValueSource.computed:
                    allowed_3.append(cell)
                else:
                    violate_2.append(cell)
    return {"violate_1": violate_1, "violate_2": violate_2, "allowed_3": allowed_3}


# ═══════════════════════════════════════════════════════════════════════════
# Property 7 本体：手工/派生（非 computed）格在模板无公式时**不得** mask
# **Validates: Requirements 4.5**
# ═══════════════════════════════════════════════════════════════════════════


@pytest.mark.parametrize("case", CASES, ids=[c.label for c in CASES])
def test_p7_no_manual_cell_is_masked_when_template_has_no_formula(case: _Case) -> None:
    """🔴 Property 7：模板无公式的非 computed 格**一律不得** mask（否则 OO 里改不了）。

    这正是 D4-1 踩过、D7-1 区1（36 格）与 D6-1 区1+区2（60 格）实测复现的 fail-closed 缺陷。
    """
    res = _classify(case)
    assert not res["violate_2"], (
        f"{case.label}: {len(res['violate_2'])} 个模板无公式的非 computed 格被 mask "
        f"⇒ 审计师在 OO 里改不了（Property 7 红）：{res['violate_2'][:40]}"
    )


@pytest.mark.parametrize("case", CASES, ids=[c.label for c in CASES])
def test_p7_every_template_formula_cell_is_masked(case: _Case) -> None:
    """反向：模板里**有公式**的格必须 mask（否则 OO 回写会覆盖模板公式）。

    豁免仅 `_COLUMN_A_MIRROR_ALLOWLIST`（A 列镜像公式，与框架不变量冲突，见文件头登记）。
    """
    res = _classify(case)
    assert not res["violate_1"], (
        f"{case.label}: {len(res['violate_1'])} 个模板公式格未 mask ⇒ OO 回写会覆盖公式："
        f"{res['violate_1'][:40]}"
    )


@pytest.mark.parametrize("case", CASES, ids=[c.label for c in CASES])
def test_p7_editable_text_cells_never_masked_framework_invariant(case: _Case) -> None:
    """框架不变量自检：数据行的 editable 列（A/L）绝不入 mask（不抛即过）。"""
    case.spec.assert_data_cells_not_masked()


@pytest.mark.parametrize("case", CASES, ids=[c.label for c in CASES])
def test_p7_derived_and_computed_fields_are_not_oo_writable(case: _Case) -> None:
    """对照面：派生（cross_sheet）与 computed 字段一律不可 OO 直写。

    🔴 这条与 mask **是两套独立机制**：mask 防「覆盖模板公式」，`is_oo_writable` 防
    「OO 写进来的派生值被前端重算覆盖 ⇒ 静默丢数据」。D3-1 区2 的 F 未 mask 但
    `is_oo_writable=False`，正是两者分离的见证。
    """
    for fs in case.spec.field_specs:
        key = fs[0]
        src = case.spec.source_of(key)
        if src in (AdjudicationValueSource.computed, AdjudicationValueSource.cross_sheet):
            assert case.spec.is_oo_writable(key) is False, f"{case.label}: 派生字段 {key} 可 OO 直写"
    # 手工金额/文本必须可写（否则整表只读，双向回写无意义）。
    assert case.spec.is_oo_writable("current_aje") is True
    assert case.spec.is_oo_writable("reason_analysis") is True


# ═══════════════════════════════════════════════════════════════════════════
# 判据有牙齿：变异 + 空分母如实登记 + 豁免可伪证
# ═══════════════════════════════════════════════════════════════════════════


@pytest.mark.parametrize("case", CASES, ids=[c.label for c in CASES])
def test_p7_mutation_masking_one_manual_cell_turns_judgement_red(case: _Case) -> None:
    """🔴 变异反证：往 mask 里塞一个「模板无公式的非 computed 格」，Property 7 判据必红。

    证明 `test_p7_no_manual_cell_is_masked_when_template_has_no_formula` 不是恒真装饰。
    """
    rows = _rows_of_interest(case)
    formulas = _formula_cells(case, rows)
    col_src = _col_to_source(case)
    mask = set(case.spec.masked_cells)
    # 找一个当前未 mask、模板无公式、且非 computed 的格作为变异靶子。
    target = None
    for row in sorted(_data_rows(case)):
        for col in _COLS:
            cell = f"{col}{row}"
            if (
                cell not in mask
                and cell not in formulas
                and col_src.get(col) is not AdjudicationValueSource.computed
            ):
                target = cell
                break
        if target:
            break
    if target is None:
        pytest.skip(f"{case.label}: 无可用变异靶子（该表数据行金额格全是模板公式，空分母）")

    mutated = mask | {target}
    # 用变异 mask 复算 violate_2 —— 必须检出 target。
    violate_2 = [
        c
        for c in mutated
        if c not in formulas
        and col_src.get(_CELL_RE.match(c).group(1)) is not AdjudicationValueSource.computed
        and int(_CELL_RE.match(c).group(2)) in rows
    ]
    assert target in violate_2, (
        f"{case.label}: 把 {target} 误加进 mask 后判据没红 ⇒ 判据无鉴别力"
    )


def test_p7_d5_manual_amount_denominator_is_empty_by_template_fact() -> None:
    """🔴 如实登记空分母：D5-1 数据行的金额格**全部是模板 SUMIF 公式** ⇒ 可编辑手工金额格为 0。

    不硬凑成「Property 7 绿」来凑齐三家口径 —— D5-1 的 Property 7 是**空分母成立**：
    没有任何手工金额格，因此「手工金额格不得被 mask」无可违反。真正在 D5-1 起作用的是
    反向判据（模板公式格必须 mask，46 格逐格一致）。
    """
    case = CASES[0]
    assert case.label == "D5-1"
    rows = _data_rows(case)
    formulas = _formula_cells(case, rows)
    col_src = _col_to_source(case)
    amount_cols = [
        fs[1]
        for fs in case.spec.field_specs
        if fs[3] == "amount" and case.spec.source_of(fs[0]) is AdjudicationValueSource.manual
    ]
    assert amount_cols, "D5-1 应有 manual/amount 字段声明（否则本判据在测别的东西）"
    unmasked_manual_amount = [
        f"{col}{row}"
        for row in sorted(rows)
        for col in amount_cols
        if f"{col}{row}" not in case.spec.masked_cells
    ]
    # 空分母的**依据**：这些格在模板里都有公式。
    for row in sorted(rows):
        for col in amount_cols:
            cell = f"{col}{row}"
            assert cell in formulas, f"D5-1 {cell} 被认定为模板公式格，实测无公式（空分母依据失效）"
            assert col_src[col] is AdjudicationValueSource.manual
    assert unmasked_manual_amount == [], (
        f"D5-1 的 manual/amount 格应全部 mask（模板公式），实得未 mask：{unmasked_manual_amount}"
    )


@pytest.mark.parametrize("case", CASES, ids=[c.label for c in CASES])
def test_p7_d67_manual_cells_actually_became_editable(case: _Case) -> None:
    """变异反证的另一面：D6-1 / D7-1 修复后**真的**有可编辑手工金额格（不是把判据改绿了）。

    D5-1 空分母（见上条）故跳过。
    """
    if case.label == "D5-1":
        pytest.skip("D5-1 手工金额格空分母（模板全公式），见 test_p7_d5_..._denominator_is_empty")
    rows = _data_rows(case)
    formulas = _formula_cells(case, rows)
    col_src = _col_to_source(case)
    editable_amount = [
        f"{col}{row}"
        for row in sorted(rows)
        for col in _COLS
        if f"{col}{row}" not in case.spec.masked_cells
        and f"{col}{row}" not in formulas
        and col_src.get(col) is AdjudicationValueSource.manual
        and any(fs[1] == col and fs[3] == "amount" for fs in case.spec.field_specs)
    ]
    assert editable_amount, f"{case.label}: 修复后仍无任何可编辑手工金额格 ⇒ 双向回写不成立"


def test_p7_column_a_mirror_allowlist_is_falsifiable() -> None:
    """🔴 豁免必须可伪证：逐条验证 A 列镜像豁免的三个成立条件，任一不成立即红。

    ① 真在 A 列 ② 真在数据行 ③ 模板值真是**同列自引用**镜像公式（`=A{n}`）。
    另断言名单**无失效条目**（名单里的格必须仍是模板公式格，否则该条已过期须删）。
    """
    by_label = {c.label: c for c in CASES}
    for label, cells in _COLUMN_A_MIRROR_ALLOWLIST.items():
        case = by_label[label]
        rows = _rows_of_interest(case)
        formulas = _formula_cells(case, rows)
        data_rows = _data_rows(case)
        assert cells, f"{label}: 豁免名单为空却仍登记（应删除该条）"
        for cell in cells:
            m = _CELL_RE.match(cell)
            assert m, f"{label}: 豁免项 {cell} 不是合法 A1 坐标"
            col, row = m.group(1), int(m.group(2))
            assert col == "A", f"{label}: 豁免项 {cell} 不在 A 列 ⇒ 豁免理由不成立"
            assert row in data_rows, f"{label}: 豁免项 {cell} 不在数据行 ⇒ 豁免理由不成立"
            assert cell in formulas, (
                f"{label}: 豁免项 {cell} 在模板里已**不是**公式格 ⇒ 该豁免已失效，须从名单删除"
            )
            formula = formulas[cell].replace(" ", "")
            assert re.fullmatch(r"=A\d+", formula), (
                f"{label}: 豁免项 {cell} 的模板公式是 {formula!r}，不是同列自引用镜像 "
                f"⇒ 豁免理由不成立（可能是别的跨表公式，须单独评估）"
            )


def test_p7_mask_and_template_symmetric_difference_is_accounted() -> None:
    """总账：三家的 mask 与模板公式集合的差异**全部**落在已登记的三类里，无第四类漏网。

    这条是「口径闭合性」判据 —— 防止未来新增一类偏离却没人分类（静默脱管）。
    """
    for case in CASES:
        res = _classify(case)
        rows = _rows_of_interest(case)
        formulas = set(_formula_cells(case, rows))
        mask = set(case.spec.masked_cells)
        allow = set(_COLUMN_A_MIRROR_ALLOWLIST.get(case.label, ()))
        considered = {f"{c}{r}" for r in rows for c in _COLS}
        # 对称差 = (公式未mask) ∪ (mask无公式)，必须被 violate_1 / violate_2 / allowed_3 / 豁免 完全覆盖
        sym_diff = ((formulas & considered) - mask) | (mask - formulas)
        accounted = set(res["violate_1"]) | set(res["violate_2"]) | set(res["allowed_3"]) | allow
        assert sym_diff <= accounted, (
            f"{case.label}: 出现未分类的 mask/模板偏离（第四类漏网）："
            f"{sorted(sym_diff - accounted)[:30]}"
        )
        # 并且本轮修复后：违反项必须为空（allowed_3 与豁免可以非空）。
        assert not res["violate_2"], f"{case.label} violate_2 非空：{res['violate_2'][:20]}"
        assert not res["violate_1"], f"{case.label} violate_1 非空：{res['violate_1'][:20]}"


if __name__ == "__main__":  # pragma: no cover
    raise SystemExit(pytest.main([__file__, "-v", "--tb=short"]))

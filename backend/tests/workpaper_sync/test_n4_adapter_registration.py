# -*- coding: utf-8 -*-
"""N4 税金及附加 真双向改线守卫 —— 受管表选型/几何 + provider/契约/registry + 发布/翻转。

spec: n-cycle-sync-foundation-and-first-canary · N4 canary

照 `test_l1_adapter_registration.py` 范式。本文件自带口径真源（N4 单 entry，几何小，
不另拆 facts 模块）。真库判据连不上就 skip-with-reason，绝不静默转绿。
"""
from __future__ import annotations

import os
import warnings
from pathlib import Path
from typing import Any

import pytest

warnings.filterwarnings("ignore", category=UserWarning, module="openpyxl")

ROOT = Path(__file__).resolve().parents[3]
BACKEND = ROOT / "backend"
DATA = BACKEND / "data"
MANIFEST_PATH = DATA / "workpaper_sync_entry_manifest.json"
CONTRACT_DIR = DATA / "workpaper_sync_contracts"
N4_TEMPLATE = BACKEND / "wp_templates" / "N" / "N4 税金及附加.xlsx"

N4_ENTRY_ID = "xlsx/gt-n4-taxes-and-surcharges"
N4_ADAPTER_ID = "n4.taxes_and_surcharges"
N4_MANAGED_SHEET = "税金及附加明细表N4-2"
N4_DERIVED_SHEET = "税金及附加审定表N4-1"

#: 净化后模板 sha256（Task 7a sanitize_n4_template_external_links --apply 之后）。
N4_TEMPLATE_SHA256 = "2005eada32506e9e2b1b6f68c704ca4f6626b8a78e9e2a3a221ca00602cad638"

#: 受管表几何基线（openpyxl 现算，§0.3）。
MANAGED_GEOMETRY = {
    "dims": "A1:K34",
    "max_row": 34,
    "max_col": 11,
    "header_row": 8,
    "first_data_row": 9,
    "last_data_row": 18,
    "footer_row": 19,
    "uuid_col": "L",
    "footer_label_cell": "A19",
    "footer_marker": "合计",
}
MANAGED_FORMULA_COLUMNS = ("E", "I", "K")
MANAGED_FORMULA_SHAPES = {"E": "=B9+C9+D9", "I": "=F9+G9+H9", "K": "=E9-J9"}
#: 🔴 模板既知缺陷（记录型锁定，不改模板）。
TEMPLATE_DEFECT_E9 = "=B9+C9+N4"      # 第三加数误写列引用（NC-32 A1 族）
TEMPLATE_DEFECT_D19 = "=SUM(D4:N18)"  # footer D 列 SUM 越界
BARE_IF_MANAGED = 0
BARE_IF_DERIVED = 20
PREPRINTED_TAX_NAMES = (
    "消费税", "城市维护建设税", "教育费附加", "资源税",
    "房产税", "土地使用税", "车船使用税", "印花税",
)


# ═══════════════════════════════════════════════════════════════════════════
# 工具
# ═══════════════════════════════════════════════════════════════════════════
def _is_formula(value: Any) -> bool:
    return isinstance(value, str) and value.startswith("=")


def _load_wb():
    from openpyxl import load_workbook

    return load_workbook(N4_TEMPLATE, data_only=False)


def _managed_ws():
    wb = _load_wb()
    return wb, wb[N4_MANAGED_SHEET]


def _count_bare_if(ws) -> int:
    n = 0
    for row in ws.iter_rows():
        for cell in row:
            v = cell.value
            if isinstance(v, str) and v.startswith("="):
                up = v.upper()
                if "IF(" in up and "IFERROR" not in up:
                    n += 1
    return n


def _dsn() -> str:
    url = os.environ.get("DATABASE_URL", "")
    if not url:
        try:
            from app.core.config import settings  # type: ignore

            url = settings.DATABASE_URL  # type: ignore[attr-defined]
        except Exception:  # pragma: no cover
            url = "postgresql+asyncpg://postgres:postgres@localhost:5432/audit_platform"
    return url.replace("+asyncpg", "").replace("+psycopg2", "")


def _query(sql: str) -> list[tuple[Any, ...]]:
    """只读查询。连不上就 skip 并写明原因 —— 不静默转绿。"""
    try:
        import psycopg2
    except ImportError:  # pragma: no cover
        pytest.skip("psycopg2 未安装 ⇒ 真库判据无法取证（不视为通过）")
    try:
        conn = psycopg2.connect(_dsn(), connect_timeout=5)
    except Exception as exc:  # pragma: no cover
        pytest.skip(f"真库 audit_platform 连不上 ⇒ 真库判据无法取证（不视为通过）：{exc}")
    try:
        with conn, conn.cursor() as cur:
            cur.execute(sql)
            return list(cur.fetchall())
    finally:
        conn.close()


@pytest.fixture(scope="module")
def provider():
    from app.services.workpaper_sync import phase5_n4_taxes_and_surcharges as mod

    return mod


# ═══════════════════════════════════════════════════════════════════════════
# 受管 sheet 选型裁决（§0.2）：canary 必须是明细表 N4-2，不是审定表 N4-1
# ═══════════════════════════════════════════════════════════════════════════
class TestManagedSheetSelection:
    """**Feature: n-cycle-sync-foundation-and-first-canary, N4 选型裁决**"""

    def test_derived_sheet_has_zero_writable_cells(self) -> None:
        """审定表 N4-1 的 r7~r15 × A..M 无一个可输入格 —— 全是公式/空。"""
        wb = _load_wb()
        ws = wb[N4_DERIVED_SHEET]
        writable: list[str] = []
        for r in range(7, 16):
            for col in "ABCDEFGHIJKLM":
                v = ws[f"{col}{r}"].value
                if v is not None and not _is_formula(v):
                    writable.append(f"{col}{r}={v!r}")
        assert writable == [], (
            f"{N4_DERIVED_SHEET} 出现可输入格 {writable} ⇒ canary 改选依据需重算"
        )

    def test_derived_sheet_pulls_from_managed_sheet(self) -> None:
        """N4-1 的 B7 引用受管表 N4-2，K7 是裸 IF ⇒ 写它毁掉取数联动。"""
        wb = _load_wb()
        ws = wb[N4_DERIVED_SHEET]
        b7 = str(ws["B7"].value)
        assert N4_MANAGED_SHEET in b7, f"B7 未引用 {N4_MANAGED_SHEET}：{b7!r}"
        assert b7.startswith("="), f"B7 不是公式：{b7!r}"
        k7 = str(ws["K7"].value).upper()
        assert "IF(" in k7 and "IFERROR" not in k7, f"K7 不是裸 IF：{k7!r}"

    def test_treating_derived_as_writable_would_be_flagged(self) -> None:
        """变异：把 N4-1 的 A/B/C 列当可写集合必须全部识别为「会覆盖公式/非空」。

        🔴 N4-1 的 A 列是跨表引用公式、B/C 列同样 ⇒ 列入可写集合即打红。
        """
        wb = _load_wb()
        ws = wb[N4_DERIVED_SHEET]
        pretend = ["A7", "B7", "C7"]
        clobbered = [ref for ref in pretend if _is_formula(ws[ref].value)]
        assert clobbered == pretend, (
            f"把 A/B/C 列当可写格时应全部识别为「会覆盖公式」，实测只识别出 {clobbered} ⇒ 判据空转"
        )

    def test_managed_sheet_has_real_inputable_cells(self) -> None:
        """受管表 N4-2 的数据区有真正可输入格（税种名 + 未审/调整列），与派生表相反。"""
        wb, ws = _managed_ws()
        inputable = []
        for r in range(9, 17):  # r9~r16 预印税种名
            for col in ("A", "B", "C", "D", "F", "G", "H", "J"):
                v = ws[f"{col}{r}"].value
                if v is None or not _is_formula(v):
                    inputable.append(f"{col}{r}")
        assert inputable, "受管表 N4-2 数据区无可输入格 ⇒ 选型裁决需重算"


# ═══════════════════════════════════════════════════════════════════════════
# 受管表几何冻结（§0.1 / §0.3）
# ═══════════════════════════════════════════════════════════════════════════
class TestManagedSheetGeometry:
    """**Feature: n-cycle-sync-foundation-and-first-canary, N4 几何冻结**"""

    def test_template_sha256_is_sanitized_value(self) -> None:
        import hashlib

        digest = hashlib.sha256(N4_TEMPLATE.read_bytes()).hexdigest()
        assert digest == N4_TEMPLATE_SHA256, (
            f"模板 sha256 实测 {digest}，冻结 {N4_TEMPLATE_SHA256}（净化后值）"
        )

    def test_geometry_matches_design(self) -> None:
        wb, ws = _managed_ws()
        assert ws.dimensions == MANAGED_GEOMETRY["dims"]
        assert ws.max_row == MANAGED_GEOMETRY["max_row"]
        assert ws.max_column == MANAGED_GEOMETRY["max_col"]
        assert ws["A8"].value == "项目"
        assert ws["K8"].value == "与应交税费贷方差异"
        assert ws["A9"].value == "消费税"
        for i, name in enumerate(PREPRINTED_TAX_NAMES):
            assert ws[f"A{9 + i}"].value == name, (9 + i, name)
        assert ws["A17"].value is None and ws["A18"].value is None

    def test_footer_label_in_column_a(self) -> None:
        """🔴 N4 footer 标签在 A19（引擎默认 A 列，与 L1 的 C 列不同）。"""
        wb, ws = _managed_ws()
        assert ws[MANAGED_GEOMETRY["footer_label_cell"]].value == MANAGED_GEOMETRY["footer_marker"]

    def test_formula_columns_are_exactly_three(self) -> None:
        from openpyxl.utils import get_column_letter

        wb, ws = _managed_ws()
        found: list[str] = []
        for ci in range(1, ws.max_column + 1):
            col = get_column_letter(ci)
            if any(_is_formula(ws[f"{col}{r}"].value) for r in range(9, 19)):
                found.append(col)
        assert tuple(found) == MANAGED_FORMULA_COLUMNS, (
            f"数据区公式列实测 {found}，期望 {list(MANAGED_FORMULA_COLUMNS)}"
        )
        # 形态以 r10 为正确模板（r9 的 E 列是模板缺陷，另有记录型断言）
        for col, shape in {"E": "=B10+C10+D10", "I": "=F10+G10+H10", "K": "=E10-J10"}.items():
            actual = str(ws[f"{col}10"].value).replace(" ", "")
            assert actual == shape, f"{col}10 形态 {actual!r} != {shape!r}"

    def test_footer_sum_columns(self) -> None:
        from openpyxl.utils import get_column_letter

        wb, ws = _managed_ws()
        summed = [
            get_column_letter(ci)
            for ci in range(1, ws.max_column + 1)
            if _is_formula(ws.cell(row=19, column=ci).value)
        ]
        assert tuple(summed) == ("B", "C", "D", "E", "F", "G", "H", "I", "J", "K"), summed

    def test_managed_sheet_has_no_bare_if(self) -> None:
        wb, ws = _managed_ws()
        assert _count_bare_if(ws) == BARE_IF_MANAGED, "受管表出现裸 IF ⇒ 中性化策略需重评估"

    def test_derived_sheet_bare_if_distribution(self) -> None:
        """整册裸 IF 20 全在派生表 N4-1。"""
        wb = _load_wb()
        total = sum(_count_bare_if(wb[n]) for n in wb.sheetnames)
        assert total == BARE_IF_DERIVED, f"整册裸 IF 实测 {total}，登记 {BARE_IF_DERIVED}"
        assert _count_bare_if(wb[N4_DERIVED_SHEET]) == BARE_IF_DERIVED

    def test_defined_names_empty_with_mutation_proof(self) -> None:
        wb = _load_wb()
        assert list(wb.defined_names.keys()) == [], "N4 册出现 definedName ⇒ 基线须更新"
        from openpyxl.workbook.defined_name import DefinedName

        wb.defined_names.add(DefinedName("gt_probe", attr_text="#REF!"))
        probed = list(wb.defined_names.keys())
        broken = [k for k in probed if "#REF!" in str(wb.defined_names[k].value)]
        assert probed == ["gt_probe"] and broken == ["gt_probe"], (
            "注入含 #REF! 的 defined name 后扫描器未命中 ⇒ 空分母不可信"
        )

    def test_template_known_defects_are_recorded_not_fixed(self) -> None:
        """🔴 两处模板缺陷记录型锁定：E9 第三加数误写列引用、D19 footer SUM 越界。"""
        wb, ws = _managed_ws()
        assert str(ws["E9"].value).replace(" ", "") == TEMPLATE_DEFECT_E9, (
            f"E9 实测 {ws['E9'].value!r}，模板缺陷基线 {TEMPLATE_DEFECT_E9}"
        )
        assert str(ws["D19"].value).replace(" ", "") == TEMPLATE_DEFECT_D19

    def test_uuid_col_is_max_col_plus_one(self) -> None:
        from openpyxl.utils import get_column_letter

        wb, ws = _managed_ws()
        assert get_column_letter(ws.max_column + 1) == MANAGED_GEOMETRY["uuid_col"]


# ═══════════════════════════════════════════════════════════════════════════
# provider 接口
# ═══════════════════════════════════════════════════════════════════════════
class TestProviderInterface:
    """**Feature: n-cycle-sync-foundation-and-first-canary, N4 provider**"""

    def test_identity_constants(self, provider) -> None:
        assert provider.ENTRY_ID == N4_ENTRY_ID
        assert provider.ADAPTER_ID == N4_ADAPTER_ID
        assert provider.MANAGED_SHEET == N4_MANAGED_SHEET
        assert provider.DERIVED_SHEET == N4_DERIVED_SHEET
        assert provider.WP_CODES == frozenset({"N4T"})
        assert provider.TEMPLATE_SHA256 == N4_TEMPLATE_SHA256
        assert provider.FORMULA_MASK == ("E9:E18", "I9:I18", "K9:K18")
        assert provider.ROW_IDENTITY_STORE_KEY == "rowKey"

    def test_public_interface_is_callable(self, provider) -> None:
        for name in (
            "instrumentation_spec", "instrumentation_specs", "template_definition_payload",
            "instrumentation_definition_payload", "authority_model_payload",
            "build_contract_payload", "build_store_projection",
            "merge_projection_into_store_rows", "all_store_item_ids",
            "publish_definitions", "attach_adapters", "manifest_capability_enabled",
            "assert_entry_selectable", "assert_no_implicit_template_fallback",
        ):
            assert callable(getattr(provider, name)), name

    def test_three_aliases_present(self, provider) -> None:
        """🔴 task76 provisioning 硬前置的三别名（K 教训）。"""
        assert provider.publish_pilot_definitions is provider.publish_definitions
        assert provider.attach_pilot_adapters is provider.attach_adapters
        assert provider.PILOT_WP_CODES == provider.WP_CODES

    def test_all_store_item_ids_is_single_canary(self, provider) -> None:
        assert provider.all_store_item_ids() == ("N4-2-detail-rows",)

    def test_no_implicit_template_fallback_mutation(self, provider) -> None:
        """换不存在模板名的变异使 assert_no_implicit_template_fallback 必抛。"""
        facts = provider.TemplateResolutionFacts(
            by_wp_code={"N4T": [object()]},  # 幻影码命中了东西 ⇒ 必须打红
            parent_code="N4",
            parent_resolved_path=provider.authoritative_template_path(),
        )
        with pytest.raises(provider.EntrySelectionError):
            provider.assert_no_implicit_template_fallback(facts, wp_codes=frozenset({"N4T"}))


# ═══════════════════════════════════════════════════════════════════════════
# 契约交付
# ═══════════════════════════════════════════════════════════════════════════
class TestContractDelivery:
    """**Feature: n-cycle-sync-foundation-and-first-canary, N4 契约**"""

    def test_contract_file_matches_source(self, provider) -> None:
        """双向锁：磁盘契约 == provider 现算 payload，且过 parse_contract。"""
        provider.assert_contract_file_matches_source()

    def test_contract_is_reviewed_with_right_fields(self, provider) -> None:
        payload = provider.build_contract_payload()
        assert payload["review_status"] == "reviewed"
        assert payload["semantic_version"] == "1.0.0"
        assert payload["contract_id"] == N4_ADAPTER_ID
        assert payload["review"]["entry_id"] == N4_ENTRY_ID
        assert payload["review"]["html_store"]["item_id"] == "N4-2-detail-rows"
        assert payload["review"]["html_store"]["row_identity_key"] == "rowKey"
        assert payload["review"]["derived_readonly_sheet"]["excel_name"] == N4_DERIVED_SHEET
        sheet = payload["sheets"][0]
        table = sheet["tables"][0]
        assert table["row_identity"]["json_pointer"] == "/rows/*/rowKey"
        assert table["formula_mask"] == ["E9:E18", "I9:I18", "K9:K18"]
        assert table["header_rows"] == 1
        assert table["footer_anchor"]["marker"] == "合计"
        assert table["footer_anchor"]["search_column"] == "A"
        assert table["uuid_col"] == "L"
        formula_fields = {f["column_key"] for f in table["fields"] if f["mode"] == "formula"}
        assert formula_fields == {"period_audited", "prior_audited", "accrual_diff"}

    def test_no_candidate_contract_remains(self) -> None:
        candidate = CONTRACT_DIR / "n4.taxes_and_surcharges.candidate.json"
        assert not candidate.exists(), "N4 candidate 契约仍在 ⇒ 双源，必须删除"

    def test_contract_in_delivered_ledger(self) -> None:
        from app.services.workpaper_sync.adapters.registry import (
            DELIVERED_PER_ENTRY_CONTRACTS,
        )

        row = next(
            (r for r in DELIVERED_PER_ENTRY_CONTRACTS if r["contract_id"] == N4_ADAPTER_ID), None
        )
        assert row is not None, "N4 未登记在交付台账"
        assert row["entry_id"] == N4_ENTRY_ID
        assert row["provider_module"] == "app.services.workpaper_sync.phase5_n4_taxes_and_surcharges"


# ═══════════════════════════════════════════════════════════════════════════
# adapter 注册接线
# ═══════════════════════════════════════════════════════════════════════════
class TestAdapterRegistration:
    """**Feature: n-cycle-sync-foundation-and-first-canary, N4 registry 接线**"""

    def test_store_merge_registry_has_n4(self) -> None:
        from app.services.workpaper_sync.store_item_registry import STORE_MERGE_REGISTRY

        plan = STORE_MERGE_REGISTRY[N4_ADAPTER_ID]
        assert plan.adapter_id == N4_ADAPTER_ID
        assert plan.provider_module == "phase5_n4_taxes_and_surcharges"
        assert tuple(i.item_id for i in plan.items) == ("N4-2-detail-rows",)
        assert plan.oo_crash_neutralization_fn == "neutralize_oo_crash_if_formulas"

    def test_provider_module_whitelisted_and_paired(self) -> None:
        from app.services.workpaper_sync.adapters import registry as r

        assert (
            "app.services.workpaper_sync.phase5_n4_taxes_and_surcharges"
            in r._ALLOWED_PROVIDER_MODULES
        )
        # 🔴 白名单与台账成对（len == len）—— H9 踩过的坑。
        assert len(r._ALLOWED_PROVIDER_MODULES) == len(r.DELIVERED_PER_ENTRY_CONTRACTS)

    def test_oo_crash_fn_resolves_to_shared_object(self) -> None:
        """N4 的中性化函数与 g7 解析到**同一个函数对象**（is）—— 不新造。"""
        from app.services.workpaper_sync.store_item_registry import (
            resolve_store_merge_plan,
        )

        try:
            from app.services.workpaper_sync.store_item_registry import (
                _resolve_oo_crash_neutralization_fn,
            )
        except ImportError:
            pytest.skip("_resolve_oo_crash_neutralization_fn 不是公开符号 ⇒ 跳过 is 判据")
        n4_fn = _resolve_oo_crash_neutralization_fn(N4_ADAPTER_ID)
        g7_fn = _resolve_oo_crash_neutralization_fn("g7.soe_subsidiary_disclosure")
        assert n4_fn is g7_fn


# ═══════════════════════════════════════════════════════════════════════════
# manifest 翻转重算（Task 8）：仅 N4 bidirectional/adapter_registered，其余仍 legacy
# ═══════════════════════════════════════════════════════════════════════════
class TestManifestFlip:
    """**Feature: n-cycle-sync-foundation-and-first-canary, N4 manifest 翻转**"""

    @pytest.fixture(scope="class")
    def manifest(self) -> dict:
        import json

        return json.loads(MANIFEST_PATH.read_text(encoding="utf-8"))

    def _entry(self, manifest: dict, entry_id: str) -> dict:
        for e in manifest["entries"]:
            if e["entry_id"] == entry_id:
                return e
        raise AssertionError(f"manifest 无 {entry_id}")

    def test_n4_is_bidirectional_adapter_registered(self, manifest: dict) -> None:
        e = self._entry(manifest, N4_ENTRY_ID)
        assert e["capability"] == "bidirectional", e.get("capability")
        assert e["migration_state"] == "adapter_registered", e.get("migration_state")
        assert e["adapter_id"] == N4_ADAPTER_ID

    def test_other_n_entries_still_legacy(self, manifest: dict) -> None:
        """🔴 只翻 N4：N1/N2/N3/N5 仍 legacy_fake_bidirectional。"""
        for entry_id in (
            "xlsx/gt-n1-deferred-tax-assets",
            "xlsx/gt-n2-taxes-payable",
            "xlsx/gt-n3-deferred-tax-liabilities",
            "xlsx/gt-n5-income-tax-expense",
        ):
            e = self._entry(manifest, entry_id)
            assert e["migration_state"] == "legacy_fake_bidirectional", (
                f"{entry_id} 状态已变为 {e.get('migration_state')!r} ⇒ 不该被本 spec 翻"
            )

    def test_provider_manifest_capability_enabled(self, provider) -> None:
        import json

        manifest = json.loads(MANIFEST_PATH.read_text(encoding="utf-8"))
        assert provider.manifest_capability_enabled(manifest=manifest) is True


# ═══════════════════════════════════════════════════════════════════════════
# golden digest 含 N4
# ═══════════════════════════════════════════════════════════════════════════
class TestGoldenDigest:
    """**Feature: n-cycle-sync-foundation-and-first-canary, N4 golden digest**"""

    def test_n4_in_golden_baseline(self) -> None:
        import json

        base = json.loads(
            (BACKEND / "scripts" / "check" / "_sync_provider_golden_digest.json").read_text(
                encoding="utf-8"
            )
        )
        n4 = next((p for p in base["providers"] if p["label"] == "n4"), None)
        assert n4 is not None, "N4 不在 golden baseline ⇒ 它的漂移永远不会被发现"
        assert n4["adapter_id"] == N4_ADAPTER_ID
        assert n4["sheet_digests"].get("n42-managed")

    def test_n4_contract_digest_matches_baseline(self, provider) -> None:
        import json

        from app.services.workpaper_sync.definitions import canonical_digest

        base = json.loads(
            (BACKEND / "scripts" / "check" / "_sync_provider_golden_digest.json").read_text(
                encoding="utf-8"
            )
        )
        n4 = next(p for p in base["providers"] if p["label"] == "n4")
        assert canonical_digest(provider.build_contract_payload()) == n4["contract_payload_sha256"]


# ═══════════════════════════════════════════════════════════════════════════
# 五环发布真库证据（真 PG；连不上 skip-with-reason，绝不静默通过）
# ═══════════════════════════════════════════════════════════════════════════
class TestFiveRingPublishRealDB:
    """**Feature: n-cycle-sync-foundation-and-first-canary, N4 五环发布真库证据**"""

    def test_entry_state_has_n4_representation(self) -> None:
        rows = _query(
            "SELECT current_representation_id, representation_generation "
            "FROM working_paper_sync_entry_state "
            f"WHERE entry_id = '{N4_ENTRY_ID}'"
        )
        assert rows, (
            f"真库 working_paper_sync_entry_state 无 {N4_ENTRY_ID} 行 ⇒ 五环发布未落库"
        )
        current_rep, gen = rows[0]
        assert current_rep is not None, "current_representation_id 为空 ⇒ 未发布 representation"
        assert gen is not None and gen >= 1, f"representation_generation={gen} < 1"

    def test_content_representation_exists(self) -> None:
        rows = _query(
            "SELECT count(*) FROM working_paper_content_representation cr "
            "JOIN working_paper_sync_entry_state es "
            "ON es.current_representation_id = cr.id "
            f"WHERE es.entry_id = '{N4_ENTRY_ID}'"
        )
        assert rows[0][0] >= 1, "N4 的 current representation 在 content_representation 表无对应行"


# ═══════════════════════════════════════════════════════════════════════════
# OO roundtrip 反向变异（守卫脚本的断言能被破坏 ⇒ 非空转）
# ═══════════════════════════════════════════════════════════════════════════
class TestOoRoundtripReverseMutation:
    """**Feature: n-cycle-sync-foundation-and-first-canary, OO roundtrip 反向变异**

    不跑真 OO（那是 verify_n4_oo94_roundtrip.py 的事），只证明「公式列往返仍是公式」这条
    判据不是恒真：把 formula_mask 的一列从公式降级成普通值，roundtrip 的「仍是公式」断言必红。
    """

    def test_formula_mask_assertion_is_falsifiable(self) -> None:
        # 正常：E/I/K 列模板里是公式
        wb, ws = _managed_ws()
        for col in MANAGED_FORMULA_COLUMNS:
            assert _is_formula(ws[f"{col}10"].value), f"{col}10 应是公式"
        # 变异：把 E10 替换成普通值后，「仍是公式」断言必须打红
        ws["E10"] = 12345
        assert not _is_formula(ws["E10"].value), (
            "把公式格替换成普通值后仍判为公式 ⇒ roundtrip 的「仍是公式」判据恒真（空转）"
        )

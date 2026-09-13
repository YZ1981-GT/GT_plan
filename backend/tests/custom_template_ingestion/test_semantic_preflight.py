"""语义 preflight 守卫。

Feature: custom-workpaper-template-ingestion-and-sync-closure
Validates: Requirements 5.1, 5.2, 5.3, 5.4, 5.5, 5.6, 5.7
"""
from __future__ import annotations

import io
from zipfile import ZipFile

from openpyxl import Workbook

from app.services.custom_template_ingestion.package_scanner import scan_package_bytes
from app.services.custom_template_ingestion.policy import (
    FINDING_EMPTY_REPORT,
    Severity,
    Verdict,
)
from app.services.custom_template_ingestion.semantic_preflight import (
    FINDING_PACKAGE_GATE,
    FINDING_SEMANTIC_COMPLETE,
    GUIDANCE_SECTION_KEYS,
    ProjectionRecommendation,
    SemanticPreflightResult,
    _finish,
    package_blocks_semantic_open,
    run_semantic_preflight,
)


def _xlsx_bytes(*, sheets: dict[str, list[tuple[str, object]]] | None = None,
                defined_names: dict[str, str] | None = None) -> bytes:
    wb = Workbook()
    # 删掉默认 Sheet，按调用方重建
    default = wb.active
    wb.remove(default)
    if not sheets:
        sheets = {"Sheet1": [("A1", 1), ("A2", "=B1+1")]}
    for name, cells in sheets.items():
        ws = wb.create_sheet(name)
        for coord, value in cells:
            ws[coord] = value
    if defined_names:
        from openpyxl.workbook.defined_name import DefinedName
        for name, attr in defined_names.items():
            wb.defined_names.add(DefinedName(name=name, attr_text=attr))
    buf = io.BytesIO()
    wb.save(buf)
    return buf.getvalue()


def test_clean_workbook_is_preflight_ready() -> None:
    payload = _xlsx_bytes(defined_names={"RowKey": "Sheet1!$A$1"})
    result = run_semantic_preflight(payload, filename_hint="ok.xlsx")
    assert result.is_valid is True
    assert result.verdict is Verdict.PREFLIGHT_READY
    assert result.semantic_opened is True
    assert result.workbook is not None
    assert result.workbook.sheet_count == 1
    assert FINDING_SEMANTIC_COMPLETE in {f.code for f in result.findings}


def test_package_security_blocker_skips_semantic_open() -> None:
    result = run_semantic_preflight(b"not-a-zip", filename_hint="bad.xlsx")
    assert result.semantic_opened is False
    assert result.is_valid is False
    assert FINDING_PACKAGE_GATE in {f.code for f in result.findings}
    assert result.projection_recommendation is ProjectionRecommendation.BLOCKED


def test_path_traversal_package_blocks_semantic() -> None:
    buf = io.BytesIO()
    with ZipFile(buf, "w") as zf:
        zf.writestr("xl/../evil.xml", "<a/>")
        zf.writestr("[Content_Types].xml", "<Types/>")
    pkg = scan_package_bytes(buf.getvalue())
    assert package_blocks_semantic_open(pkg) is True


def test_workbook_report_has_digests_and_sheets() -> None:
    payload = _xlsx_bytes(sheets={
        "A": [("A1", "title"), ("B1", "=Sheet2!A1")],
        "Sheet2": [("A1", 10)],
    })
    result = run_semantic_preflight(payload)
    assert result.workbook is not None
    assert len(result.workbook.artifact_sha256) == 64
    assert result.workbook.policy_version
    assert result.workbook.scanner_build_id
    assert result.workbook.sheet_count == 2
    assert any("Sheet2!" in f or "Sheet2" in f for f in result.workbook.cross_sheet_formulas) or any(
        s.cross_sheet_formula_count > 0 for s in result.workbook.sheets
    )


def test_sheet_report_covers_required_observations() -> None:
    payload = _xlsx_bytes(sheets={"Main": [
        ("A1", "hdr"), ("A2", 1), ("B2", "=A2*2"),
    ]}, defined_names={"MainKey": "Main!$A$2"})
    result = run_semantic_preflight(payload)
    sheet = result.workbook.sheets[0]
    assert sheet.sheet_uid.startswith("sheet:")
    assert sheet.sheet_name == "Main"
    assert sheet.order == 0
    assert sheet.non_empty_cells >= 3
    assert sheet.formula_count >= 1
    assert sheet.used_range
    assert sheet.projection_recommendation in set(ProjectionRecommendation)


def test_guidance_candidates_are_nine_sections_review_pending() -> None:
    payload = _xlsx_bytes()
    result = run_semantic_preflight(payload)
    sheet = result.workbook.sheets[0]
    keys = [g.key for g in sheet.guidance_candidates]
    assert keys == list(GUIDANCE_SECTION_KEYS)
    for g in sheet.guidance_candidates:
        d = g.to_dict()
        assert d["extractionSource"] == "custom_candidate"
        assert d["completionStatus"] == "review_pending"
        assert "confirmed" not in d


def test_no_carrier_recommends_read_only_or_oo_not_editable() -> None:
    payload = _xlsx_bytes(sheets={"Bare": [("A1", "x"), ("A2", 1)]})
    result = run_semantic_preflight(payload)
    sheet = result.workbook.sheets[0]
    assert sheet.projection_recommendation in {
        ProjectionRecommendation.READ_ONLY_HTML,
        ProjectionRecommendation.ONLYOFFICE_ONLY,
    }
    assert sheet.projection_recommendation is not ProjectionRecommendation.EDITABLE_GRID
    assert "no_existing_stable_carrier" in sheet.identity_risks


def test_defined_name_carrier_allows_editable_grid_recommendation() -> None:
    payload = _xlsx_bytes(
        sheets={"Main": [("A1", 1), ("A2", 2)]},
        defined_names={"StableRow": "Main!$A$1"},
    )
    result = run_semantic_preflight(payload)
    sheet = result.workbook.sheets[0]
    assert any(c.kind.value == "defined_name" for c in sheet.identity_carriers)
    assert sheet.projection_recommendation is ProjectionRecommendation.EDITABLE_GRID


def test_empty_findings_path_never_valid() -> None:
    pkg = scan_package_bytes(_xlsx_bytes(), filename_hint="ok.xlsx")
    # 人为造一个「空 findings」收尾路径
    result = _finish(pkg, None, [], ProjectionRecommendation.EDITABLE_GRID, opened=False)
    assert result.is_valid is False
    assert any(f.code == FINDING_EMPTY_REPORT for f in result.findings)
    assert result.projection_recommendation is ProjectionRecommendation.BLOCKED


def test_formulas_not_evaluated() -> None:
    """公式格保留表达式文本，不出现求值后的数字替代。"""
    payload = _xlsx_bytes(sheets={"S": [("A1", 2), ("A2", "=A1*10")]})
    result = run_semantic_preflight(payload)
    samples = result.workbook.sheets[0].sample_formulas
    assert any(s.startswith("=A1") for s in samples)


def test_finding_shape_has_code_severity_locator_remediation() -> None:
    result = run_semantic_preflight(b"", filename_hint="x.xlsx")
    assert result.findings
    for f in result.findings:
        d = f.to_dict()
        assert "code" in d and "severity" in d
        assert "locator" in d and "remediation" in d
        assert "policyDecision" in d


def test_vba_feature_still_opens_semantic_but_not_valid() -> None:
    """能力矩阵 BLOCKER 不等于安全门：仍可读结构，但 overall 不可 valid。"""
    # 最小 xlsx + vbaProject.bin 部件
    base = _xlsx_bytes()
    buf = io.BytesIO()
    with ZipFile(io.BytesIO(base), "r") as src, ZipFile(buf, "w") as dst:
        for info in src.infolist():
            dst.writestr(info, src.read(info.filename))
        dst.writestr("xl/vbaProject.bin", b"MZ")
    result = run_semantic_preflight(buf.getvalue(), filename_hint="m.xlsx")
    # 包门未拦路径/加密 → 可打开语义；但 VBA BLOCKER 使 verdict BLOCKED
    assert result.semantic_opened is True
    assert result.is_valid is False
    assert result.projection_recommendation is ProjectionRecommendation.BLOCKED


def test_to_dict_roundtrip_keys() -> None:
    result = run_semantic_preflight(_xlsx_bytes(defined_names={"K": "Sheet1!$A$1"}))
    d = result.to_dict()
    assert d["valid"] is True
    assert d["workbook"]["sheetCount"] == 1
    assert isinstance(d["findings"], list)

"""D4 sheet label 映射 ↔ 权威模板册 对账守卫。

背景（2026-09-28）：`d4SheetLabels.ts` 的 `D4_SHEET_LABEL_MAP` 把
`D4-附注国企` 写成 `附注披露信息（国有企业）`，而三本权威册里的真名是
`附注披露信息（国企）`。因为 `resolveD4SheetLabel` 有
``includes('附注') && (includes('国企') || includes('国有'))`` 兜底，
**错名长期不可见**（fallback 每次都命中真 sheet ⇒「错配置 + 正确结果」）。
这与 M 轮 MC-23 记录的「错 tab 非空白」是同一型缺陷。

本守卫把「label 值必须是权威册里真实存在的 sheet 名」变成机器判据，
使这类错名在 CI 即暴露，而不是等到某天 fallback 失效才发现。

分母 = `backend/wp_templates/D/D4*.xlsx` 十本册 sheetnames 的**并集**（宽松口径：
某 label 只要在任一 D4 册内存在即合法，避免因分册裁剪误报）。

判据是**内容**（名字逐字存在），不是「文件存在」。
"""

import re
from pathlib import Path

import pytest
from openpyxl import load_workbook

_REPO_ROOT = Path(__file__).resolve().parents[2]
_LABELS_TS = (
    _REPO_ROOT
    / "audit-platform"
    / "frontend"
    / "src"
    / "components"
    / "workpaper"
    / "composables"
    / "d4SheetLabels.ts"
)
_TEMPLATE_DIR = _REPO_ROOT / "backend" / "wp_templates" / "D"

# 哨兵值：不是 sheet 名，而是宿主 `useD4EntryDualMode.resolveOoSheetName` 的特判返回
#   （`code === 'D4' || 'skip' || 'directory'` → return 'D4'），故不参与册内存在性对账。
_SENTINEL_VALUES = {"D4"}


def _read_text(path: Path) -> str:
    raw = path.read_bytes()
    if raw.startswith(b"\xef\xbb\xbf"):
        raw = raw[3:]
    return raw.decode("utf-8")


def _parse_label_map(src: str) -> dict[str, str]:
    """解析 `D4_SHEET_LABEL_MAP` 的键值对（键可带引号也可裸写，如 `D4A: '...'`）。"""
    block = re.search(
        r"D4_SHEET_LABEL_MAP\s*:\s*Record<string,\s*string>\s*=\s*\{(.*?)\n\}",
        src,
        re.S,
    )
    assert block, "未能在 d4SheetLabels.ts 中定位 D4_SHEET_LABEL_MAP 定义块"
    body = block.group(1)
    # 逐行取 `key: 'value',`，先剥掉 // 注释行避免把注释里的示例名当配置
    pairs: dict[str, str] = {}
    for line in body.split("\n"):
        stripped = line.strip()
        if not stripped or stripped.startswith("//"):
            continue
        m = re.match(r"^'([^']+)'\s*:\s*'([^']*)'\s*,?$", stripped) or re.match(
            r"^([A-Za-z0-9_$]+)\s*:\s*'([^']*)'\s*,?$", stripped
        )
        if m:
            pairs[m.group(1)] = m.group(2)
    return pairs


def _workbook_sheet_name_union() -> set[str]:
    names: set[str] = set()
    for book in sorted(_TEMPLATE_DIR.glob("D4*.xlsx")):
        wb = load_workbook(book, read_only=True)
        try:
            names.update(wb.sheetnames)
        finally:
            wb.close()
    return names


@pytest.fixture(scope="module")
def label_map() -> dict[str, str]:
    assert _LABELS_TS.exists(), f"{_LABELS_TS} 不存在"
    return _parse_label_map(_read_text(_LABELS_TS))


@pytest.fixture(scope="module")
def sheet_names() -> set[str]:
    return _workbook_sheet_name_union()


def test_parser_self_check(label_map: dict[str, str]):
    """守卫自检（防恒真）：解析器必须真的取到足量条目与已知样本。"""
    assert len(label_map) >= 40, f"仅解析出 {len(label_map)} 条，解析器可能失配"
    # 带引号键 + 裸键 两种写法都要能取到
    assert label_map.get("D4-1") == "营业收入审定表D4-1"
    assert label_map.get("D4A") == "营业收入审计程序表D4A"


def test_template_dir_has_d4_books(sheet_names: set[str]):
    """分母自检：权威册目录必须真的读出 sheet 名（否则下方断言恒真）。"""
    assert len(sheet_names) >= 40, f"D4 册 sheetnames 并集仅 {len(sheet_names)} 条"
    assert "营业收入调整分录汇总D4-4" in sheet_names


def test_every_label_exists_in_authoritative_workbook(
    label_map: dict[str, str], sheet_names: set[str]
):
    """每个 label 值必须是权威册里逐字存在的 sheet 名（哨兵值除外）。

    这条断言直接封死「错名被 resolveD4SheetLabel 的 includes 兜底掩盖」的整类缺陷。
    """
    missing = {
        code: name
        for code, name in label_map.items()
        if name not in _SENTINEL_VALUES and name not in sheet_names
    }
    assert not missing, (
        "以下 label 的值在 D4 权威册里不存在（错名 / 过期名）：\n"
        + "\n".join(f"  {code!r} -> {name!r}" for code, name in sorted(missing.items()))
    )


def test_note_soe_label_is_exact_template_name(label_map: dict[str, str]):
    """单独钉死本次修复的那条：国企附注真名是「（国企）」不是「（国有企业）」。"""
    assert label_map["D4-附注国企"] == "附注披露信息（国企）"
    assert label_map["D4-附注上市"] == "附注披露信息（上市公司）"


def test_d4_4_label_is_exact_template_name(label_map: dict[str, str], sheet_names: set[str]):
    """D4-4 的 label 与权威册逐字一致（legacy 禁入后仍需正确解析，勿回归）。"""
    assert label_map["D4-4"] == "营业收入调整分录汇总D4-4"
    assert label_map["D4-4"] in sheet_names


def test_sentinel_values_are_declared_not_accidental(label_map: dict[str, str]):
    """哨兵值必须仅限显式登记的那些，防止新增错名被当哨兵放过。"""
    sentinel_codes = sorted(c for c, n in label_map.items() if n in _SENTINEL_VALUES)
    assert sentinel_codes == ["D4", "D4-目录"], (
        f"哨兵值使用范围发生变化：{sentinel_codes}；"
        "新增前请确认它确实不是 sheet 名（对应 resolveOoSheetName 的特判）"
    )

"""服务端完整导出：导出闸、文本强制与响应收尾（advanced-query-linkage-closure 任务 26；R8.2、R8.5、R8.6；design §3.5、§10.1）。

- 导出闸（``export_gate``）：同一用户第 2 个 429、全进程第 5 个 429、释放后可再取、``release`` 幂等（旧名额重复释放不放行新名额）；
- 文本强制（``export_service._text_safe_cell``）：以 ``= + - @`` 开头的文本在数据、表头、「查询说明」里都按文本写入，数值 −1
  不受影响（变异 #18）；
- 构建器端点（SQLite，只建构建器白名单表，同 ``test_query_builder_streaming_export.py``）：客户端中途断开（ASGI 两种形态：
  先断开后发送 / 发送时连接已断）后导出名额与临时文件都被释放、同一用户可立即再导（变异 B7-a、B7-d）；COUNT 超过上限 400、
  不开始取数、不产生文件（变异 B7-b）；本人已有在途导出 429；表头取请求的 ``headers``、附「查询说明」、RFC 5987 文件名、记审计。
- sheet 名清洗（``export_service.safe_sheet_title``，B7 评审 B7-R1）：Excel 不允许的 ``\\ / * ? : [ ]`` 换成全角同形字符、去首尾
  单引号、截 31 字符、空则缺省名；白名单每个表名清洗后 openpyxl 都接受；构建器导出「调整分录（AJE/RJE）」200（改前
  ``create_sheet`` 抛 ValueError、恒 500；变异 B7-f / B7-f2 / B7-f3）。
- 枚举单元格（B7 复测发现）：构建器按 ORM 列取数，``sa.Enum`` 列是枚举成员；导出写枚举值（aje / draft），不写「类名.成员」
  （变异 B7-g）。
"""
from __future__ import annotations

import asyncio
import io
import json
import tempfile
import uuid
from datetime import date
from decimal import Decimal
from urllib.parse import quote

import pytest
import pytest_asyncio
from fastapi import FastAPI, HTTPException
from httpx import ASGITransport, AsyncClient
from openpyxl import Workbook, load_workbook
from openpyxl.cell.cell import Cell
from openpyxl.workbook.child import INVALID_TITLE_REGEX
from sqlalchemy.dialects.sqlite.base import SQLiteTypeCompiler
from sqlalchemy.ext.asyncio import AsyncSession, async_sessionmaker, create_async_engine

from app.core.database import get_db
from app.deps import get_current_user
from app.models.audit_platform_models import AccountCategory, Adjustment, AdjustmentType, ReviewStatus, TrialBalance
from app.models.core import Project, ProjectStatus, ProjectType, UserRole
from app.routers.query_builder import router as query_builder_router
from app.services.custom_query import audit_helper, export_gate, paged_fetch
from app.services.custom_query import export_service as export_module
from app.services.custom_query.export_service import ExportService, _excel_cell_value, _text_safe_cell, safe_sheet_title
from app.services.custom_query.query_orchestrator import ColumnMeta
from app.services.custom_query.table_whitelist import TABLE_WHITELIST
from tests.advanced_query_linkage.builder_tables import assert_only_builder_tables, create_builder_schema

SQLiteTypeCompiler.visit_JSONB = SQLiteTypeCompiler.visit_JSON

EXPORT = "/api/query/export-excel"
FIELDS = ["standard_account_code", "account_name", "audited_amount"]
HEADERS = {"standard_account_code": "标准科目编码", "account_name": "科目名称", "audited_amount": "审定数"}
USER_ID = uuid.uuid4()


class _Admin:
    id = USER_ID
    role = UserRole.admin
    username = "aql-b7-admin"
    email = "aql-b7@test.invalid"
    is_active = True


@pytest.fixture(autouse=True)
def _gate_is_empty():
    """每例前后导出闸都必须是空的：名额泄漏会让后续用例碰巧 429，在这里当场暴露。"""
    assert export_gate.in_flight_count() == 0, "上一个用例泄漏了导出名额"
    yield
    leaked = export_gate.in_flight_count()
    export_gate._IN_FLIGHT.clear()
    assert leaked == 0, f"本例结束时仍有 {leaked} 个导出名额未释放"


# ─── 导出闸 ───────────────────────────────────────────────────────────────────


def test_same_user_second_export_is_busy():
    slot = export_gate.acquire_export_slot("u1")
    with pytest.raises(HTTPException) as exc:
        export_gate.acquire_export_slot("u1")
    assert exc.value.status_code == 429
    assert exc.value.detail == {"error_code": "EXPORT_BUSY", "message": "已有导出在进行，请稍后再试"}
    slot.release()
    export_gate.acquire_export_slot("u1").release()


def test_fifth_concurrent_export_in_the_process_is_busy():
    slots = [export_gate.acquire_export_slot(f"u{i}") for i in range(export_gate.MAX_EXPORTS_PER_PROCESS)]
    assert export_gate.MAX_EXPORTS_PER_PROCESS == 4
    with pytest.raises(HTTPException) as exc:
        export_gate.acquire_export_slot("u-fifth")
    assert exc.value.status_code == 429 and exc.value.detail["error_code"] == "EXPORT_BUSY"
    slots[0].release()
    fifth = export_gate.acquire_export_slot("u-fifth")
    for slot in (*slots[1:], fifth):
        slot.release()


def test_release_is_idempotent_and_a_stale_slot_cannot_free_a_new_one():
    first = export_gate.acquire_export_slot("u1")
    first.release()
    first.release()
    assert export_gate.in_flight_count() == 0 and first.released
    second = export_gate.acquire_export_slot("u1")
    first.release()  # 旧名额再释放一次：不得把同一用户新取的名额放掉
    with pytest.raises(HTTPException):
        export_gate.acquire_export_slot("u1")
    second.release()
    assert export_gate.in_flight_count() == 0


# ─── 文本强制 ─────────────────────────────────────────────────────────────────


@pytest.mark.parametrize("text", ["=1+1", "+1", "-1", "@x"])
def test_formula_like_text_becomes_a_text_cell(text: str):
    ws = Workbook(write_only=True).create_sheet()
    cell = _text_safe_cell(ws, text)
    assert isinstance(cell, Cell) and cell.data_type == "s" and cell.value == text  # WriteOnlyCell 是工厂函数，返回 Cell


@pytest.mark.parametrize("value", [-1, Decimal("-1.5"), 0, "摘要", "a=1", None, date(2025, 1, 1)])
def test_other_values_are_left_alone(value):
    ws = Workbook(write_only=True).create_sheet()
    assert _text_safe_cell(ws, value) is value


@pytest.mark.parametrize(("member", "value"), [
    (AdjustmentType.aje, "aje"), (ReviewStatus.draft, "draft"), (AccountCategory.asset, "asset")])
def test_enum_members_are_exported_as_their_value(member, value: str):
    got = _excel_cell_value(member)
    # (str, Enum) 成员本身就 == 其值：只比 == 看不出问题，要看类型是不是纯 str
    assert type(got) is str and got == value


async def test_stream_xlsx_writes_enum_values_not_member_names():
    content = await _collect(ExportService().stream_xlsx(
        [{"t": AdjustmentType.aje, "s": ReviewStatus.draft}], [ColumnMeta("t", "类型"), ColumnMeta("s", "状态")]))
    wb = load_workbook(io.BytesIO(content), read_only=True)
    try:
        rows = list(wb.worksheets[0].iter_rows(values_only=True))
    finally:
        wb.close()
    assert rows == [("类型", "状态"), ("aje", "draft")]


async def _collect(agen) -> bytes:
    return b"".join([chunk async for chunk in agen])


async def test_stream_xlsx_writes_formula_like_text_as_text_in_data_header_and_info():
    columns = [ColumnMeta(key="a", title="=表头"), ColumnMeta(key="b", title="金额")]
    rows = [{"a": "=1+1", "b": -1}, {"a": "+1", "b": 2}, {"a": "@x", "b": None}, {"a": "-1", "b": 3}]
    content = await _collect(ExportService().stream_xlsx(rows, columns, info_rows=[("条件", "=SUM(A1)"), ("行数", 4)]))
    wb = load_workbook(io.BytesIO(content))  # 非只读：要看单元格类型
    data, info = wb.worksheets
    grid = [[(c.value, c.data_type) for c in row] for row in data.iter_rows()]
    assert grid[0] == [("=表头", "s"), ("金额", "s")]
    assert [row[0] for row in grid[1:]] == [("=1+1", "s"), ("+1", "s"), ("@x", "s"), ("-1", "s")]
    assert [row[1] for row in grid[1:]] == [(-1, "n"), (2, "n"), (None, "n"), (3, "n")]
    assert [[(c.value, c.data_type) for c in row] for row in info.iter_rows()] == [
        [("条件", "s"), ("=SUM(A1)", "s")], [("行数", "s"), (4, "n")]]
    assert info.title == "查询说明"


async def test_info_rows_callable_receives_the_written_row_count():
    seen: list[int] = []

    def info(written: int):
        seen.append(written)
        return [("行数", written)]

    content = await _collect(ExportService().stream_xlsx([{"a": 1}, [{"a": 2}, {"a": 3}]], [ColumnMeta("a", "A")],
                                                        info_rows=info))
    assert seen == [3]
    notes = {r[0]: r[1] for r in load_workbook(io.BytesIO(content)).worksheets[1].iter_rows(values_only=True)}
    assert notes == {"行数": 3}


# ─── sheet 名清洗（B7 评审 B7-R1）──────────────────────────────────────────────


@pytest.mark.parametrize(("raw", "expected"), [
    ("调整分录（AJE/RJE）", "调整分录（AJE／RJE）"),
    ("a\\b*c?d:e[f]g/h", "a＼b＊c？d：e［f］g／h"),
    ("'引号'", "引号"),
    ("x" * 40, "x" * 31),
    ("x" * 30 + "'y", "x" * 30),  # 截到 31 个字符后落在单引号上：结尾的单引号再去掉
    (None, "高级查询结果"),
    ("", "高级查询结果"),
    ("  ' '  ", "高级查询结果"),
])
def test_safe_sheet_title_gives_a_title_excel_accepts(raw, expected: str):
    title = safe_sheet_title(raw)
    assert title == expected
    assert Workbook(write_only=True).create_sheet(title=title).title == title


def test_every_builder_table_label_becomes_a_valid_sheet_title():
    """构建器导出以白名单表名作数据 sheet 名：每个表名清洗后 openpyxl 都接受；本来合法的表名原样保留。"""
    for table, meta in TABLE_WHITELIST.items():
        label = str(meta["label"])
        title = safe_sheet_title(label)
        assert Workbook(write_only=True).create_sheet(title=title).title == title, table
        assert len(title) <= 31 and not INVALID_TITLE_REGEX.search(title), (table, title)
        if not INVALID_TITLE_REGEX.search(label) and len(label) <= 31 and label == label.strip(" '"):
            assert title == label, table
    assert safe_sheet_title(TABLE_WHITELIST["adjustments"]["label"]) == "调整分录（AJE／RJE）"


async def test_stream_xlsx_cleans_the_data_and_extra_sheet_titles():
    content = await _collect(ExportService().stream_xlsx(
        [{"a": 1}], [ColumnMeta("a", "A")], sheet_title="调整分录（AJE/RJE）", info_rows=[("行数", 1)],
        extra_sheets=[("取值:A1/B2", [["x"]])]))
    wb = load_workbook(io.BytesIO(content), read_only=True)
    try:
        assert wb.sheetnames == ["调整分录（AJE／RJE）", "查询说明", "取值：A1／B2"]
    finally:
        wb.close()


# ─── 构建器端点（SQLite）────────────────────────────────────────────────────


@pytest_asyncio.fixture
async def builder_app(monkeypatch, tmp_path):
    """构建器 router + SQLite 会话（3 行试算表、1 行调整分录）+ 系统管理员；临时文件目录与审计都换成本例私有的。"""
    xlsx_tmp = tmp_path / "xlsx_tmp"
    xlsx_tmp.mkdir()
    monkeypatch.setattr(tempfile, "tempdir", str(xlsx_tmp))
    audits: list[dict] = []

    async def _record(**kwargs):  # 审计写入的下一层：只记录，不进共享 Redis
        audits.append(kwargs)
        return True

    monkeypatch.setattr(audit_helper, "_log_unthrottled", _record)
    engine = create_async_engine("sqlite+aiosqlite:///:memory:", echo=False)
    async with engine.begin() as conn:
        await conn.run_sync(create_builder_schema)
        await assert_only_builder_tables(conn)
    factory = async_sessionmaker(engine, class_=AsyncSession, expire_on_commit=False)
    async with factory() as session:
        project_id = uuid.uuid4()
        session.add(Project(id=project_id, name="【测试】AQL-B7 导出", client_name="T", project_type=ProjectType.annual,
                            status=ProjectStatus.planning, created_by=USER_ID, audit_period_start=date(2025, 1, 1),
                            audit_period_end=date(2025, 12, 31)))
        await session.flush()
        for i, code in enumerate(("1001", "1002", "1122")):
            session.add(TrialBalance(id=uuid.uuid4(), project_id=project_id, year=2025, company_code="001",
                                     standard_account_code=code, account_name=f"科目{i}",
                                     account_category=AccountCategory.asset, audited_amount=Decimal(f"{i}00.50")))
        session.add(Adjustment(id=uuid.uuid4(), project_id=project_id, year=2025, company_code="001",
                               adjustment_no="AJE-1", adjustment_type=AdjustmentType.aje, account_code="1001",
                               debit_amount=Decimal("100.00"), entry_group_id=uuid.uuid4(), created_by=USER_ID))
        await session.flush()
        app = FastAPI()
        app.include_router(query_builder_router)

        async def _db():
            yield session

        async def _user():
            return _Admin()

        app.dependency_overrides[get_db] = _db
        app.dependency_overrides[get_current_user] = _user
        app.state.audits = audits
        app.state.xlsx_tmp = xlsx_tmp
        yield app
    await engine.dispose()


def _payload(**extra) -> dict:
    return {"table": "trial_balance", "fields": FIELDS, "limit": 1, **extra}


async def _post(app: FastAPI, payload: dict):
    async with AsyncClient(transport=ASGITransport(app=app), base_url="http://aql-b7") as client:
        return await client.post(EXPORT, json=payload)


def _leftover_files(app: FastAPI) -> list[str]:
    return sorted(p.name for p in app.state.xlsx_tmp.glob("aqm_export_*"))


async def test_export_has_chinese_headers_info_sheet_rfc5987_name_and_audit(builder_app):
    response = await _post(builder_app, _payload(headers=HEADERS, file_name="试算表/导出:全部"))
    assert response.status_code == 200, response.text[:500]
    wb = load_workbook(io.BytesIO(response.content), read_only=True)
    data, info = wb.worksheets
    rows = list(data.iter_rows(values_only=True))
    assert rows[0] == tuple(HEADERS[f] for f in FIELDS)
    assert len(rows) == 4  # limit=1 不参与导出：3 行全在
    assert data.title == "试算表"
    notes = {r[0]: r[1] for r in info.iter_rows(values_only=True)}
    assert list(notes)[:7] == ["项目", "年度", "数据源", "条件", "行数", "导出时间", "导出人"]
    assert notes["行数"] == 3 and notes["导出人"] == "aql-b7-admin"
    assert notes["数据源"] == "高级构建器：试算表（trial_balance）"
    wb.close()
    disposition = response.headers["content-disposition"]
    assert "filename*=UTF-8''" + quote("试算表_导出_全部.xlsx", safe="") in disposition
    audits = builder_app.state.audits
    assert [a["action"] for a in audits] == ["custom_query.export"]
    assert audits[0]["details"]["entry"] == "builder" and audits[0]["details"]["row_count"] == 3
    assert _leftover_files(builder_app) == []


async def test_adjustments_export_uses_a_cleaned_sheet_title(builder_app):
    """表名「调整分录（AJE/RJE）」含 '/'：数据 sheet 名清洗为全角「／」、导出 200（B7 评审 B7-R1：改前 openpyxl
    ``create_sheet`` 抛 ``ValueError: Invalid character / found in sheet title``，导出恒 500）。「查询说明」里照写原表名。"""
    fields = ["adjustment_no", "adjustment_type", "account_code", "debit_amount", "review_status"]
    response = await _post(builder_app, {"table": "adjustments", "fields": fields})
    assert response.status_code == 200, response.text[:500]
    wb = load_workbook(io.BytesIO(response.content), read_only=True)
    try:
        assert wb.sheetnames == ["调整分录（AJE／RJE）", "查询说明"]
        rows = list(wb.worksheets[0].iter_rows(values_only=True))
        notes = {r[0]: r[1] for r in wb.worksheets[1].iter_rows(values_only=True)}
    finally:
        wb.close()
    # 枚举列写枚举值（与 execute 返回、界面显示一致），不是「AdjustmentType.aje」
    assert rows == [tuple(fields), ("AJE-1", "aje", "1001", 100, "draft")]
    assert notes["数据源"] == "高级构建器：调整分录（AJE/RJE）（adjustments）" and notes["行数"] == 1
    assert [a["action"] for a in builder_app.state.audits] == ["custom_query.export"]
    assert _leftover_files(builder_app) == []


async def test_row_count_over_the_limit_is_rejected_before_any_row_is_read(builder_app, monkeypatch):
    monkeypatch.setattr(export_module, "EXPORT_ROW_HARD_LIMIT", 2)
    streamed: list[tuple] = []
    real_stream = paged_fetch.stream_statement

    def _spy(*args, **kwargs):
        streamed.append(args)
        return real_stream(*args, **kwargs)

    monkeypatch.setattr(paged_fetch, "stream_statement", _spy)
    response = await _post(builder_app, _payload())
    assert response.status_code == 400, f"超过导出上限应 400 且不产生文件，实得 {response.status_code}"
    detail = response.json()["detail"]
    assert detail["error_code"] == "EXPORT_ROW_HARD_LIMIT"
    assert (detail["actual_rows"], detail["row_hard_limit"]) == (3, 2)
    assert streamed == [], "超过导出上限时开始了取数（COUNT 预检未在取数之前拦截）"
    assert builder_app.state.audits == [] and _leftover_files(builder_app) == []


async def test_same_user_with_an_export_in_flight_gets_429(builder_app):
    held = export_gate.acquire_export_slot(USER_ID)
    try:
        response = await _post(builder_app, _payload())
        assert response.status_code == 429
        assert response.json()["detail"] == {"error_code": "EXPORT_BUSY", "message": "已有导出在进行，请稍后再试"}
    finally:
        held.release()
    assert (await _post(builder_app, _payload())).status_code == 200


async def _post_then_disconnect(app: FastAPI, payload: dict, *, spec_version: str) -> dict:
    """直接驱动 ASGI：请求体之后客户端断开。

    ``spec_version`` < 2.4：Starlette 并发监听断开，``receive`` 立即给 ``http.disconnect``；写响应头时让出一次事件循环（真实
    传输层写头同样可能让出），取消就落在这里——字节流一次都还没迭代（``chunks == 0``）。≥ 2.4：Starlette 只靠 ``send`` 报错
    发现断开，这里在第一个数据块上抛 ``OSError``（连接已断，``chunks == 1``）。
    """
    body = json.dumps(payload).encode("utf-8")
    seen: dict = {"status": None, "chunks": 0, "error": None}
    delivered = False

    async def receive() -> dict:
        nonlocal delivered
        if not delivered:
            delivered = True
            return {"type": "http.request", "body": body, "more_body": False}
        return {"type": "http.disconnect"}

    async def send(message: dict) -> None:
        if message["type"] == "http.response.start":
            seen["status"] = message["status"]
            await asyncio.sleep(0)
        elif message["type"] == "http.response.body" and message.get("body"):
            seen["chunks"] += 1
            if spec_version >= "2.4":
                raise OSError("客户端已断开")

    scope = {
        "type": "http", "asgi": {"version": "3.0", "spec_version": spec_version}, "http_version": "1.1",
        "method": "POST", "scheme": "http", "path": EXPORT, "raw_path": EXPORT.encode(), "root_path": "",
        "query_string": b"", "client": ("127.0.0.1", 50000), "server": ("aql-b7", 80),
        "headers": [(b"content-type", b"application/json"), (b"content-length", str(len(body)).encode())],
    }
    try:
        await app(scope, receive, send)
    except Exception as exc:  # noqa: BLE001 — 断开以异常冒泡（ClientDisconnect）属正常形态
        seen["error"] = type(exc).__name__
    return seen


@pytest.mark.parametrize("spec_version", ["2.3", "2.4"])
async def test_client_disconnect_releases_the_slot_and_the_temp_file(builder_app, spec_version: str):
    seen = await _post_then_disconnect(builder_app, _payload(), spec_version=spec_version)
    # 断开的位置：< 2.4 在开始发送数据之前（字节流从未启动），≥ 2.4 在第一个数据块上
    assert (seen["status"], seen["chunks"]) == (200, 0 if spec_version < "2.4" else 1), seen
    assert export_gate.in_flight_count() == 0, f"客户端断开（ASGI {spec_version}）后导出名额未释放：{seen}"
    assert _leftover_files(builder_app) == [], "客户端断开后临时文件仍在磁盘上"
    again = await _post(builder_app, _payload())
    assert again.status_code == 200, f"断开后同一用户再导出应 200，实得 {again.status_code}：{again.text[:200]}"

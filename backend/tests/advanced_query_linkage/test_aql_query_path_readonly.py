"""查询路径只读（advanced-query-linkage-closure 任务 16；R4.4；design §3.7；变异 #10、#10b）。

一、AST 守卫（``services/custom_query/fetchers/workpaper_cells.py``，判定用 AST，不做文本匹配）：
- 模块不引用 ``init_workpaper_from_template`` / ``get_workpaper_storage_path``（前者从模板写文件，后者 ``mkdir``）；
  ``routers/custom_query.py`` 也不再直接碰这些 storage 辅助（wp-sheet-preview 改调本模块）；
- 写盘调用（``shutil.copy*`` / ``move`` 等、``os.replace`` / ``rename`` / ``remove`` 等、``.write_bytes`` / ``.write_text`` /
  ``.save`` / ``.mkdir`` 等、写模式 ``open``）只出现在 ``_temp_copy`` 里；
- ``_temp_copy`` 的每个调用点第二个实参都是 ``with tempfile.TemporaryDirectory() as <名字>`` 绑定的目录（可包 ``Path``）；
- ``recompute_with_libreoffice`` 标注返回 ``bytes | None``、只返回 ``None`` 或 ``….read_bytes()``；它起的 LibreOffice 子进程
  ``--outdir`` 是临时目录、输入是 ``_temp_copy`` 的产物（不是原文件）。
扫描器自带双向自检：每种写盘形态都认得，``str.replace`` / 只读 ``open`` / ``dict.copy`` 等不误认。

二、行为（真 PG 权限世界 + 底稿存储根换成本例临时目录，请求经生产注册的 custom_query router 真发）：
有文件 ⇒ 预览 / 区域取数前后 storage 的文件集合、字节（sha256）与 mtime 都不变，LibreOffice 只在临时副本上跑（替身记录参数）；
无文件 ⇒ 预览 ``available: false`` + 「该底稿尚未生成文件…」、区域取数 400 同文案，storage 里什么都不留（目录也不建）；
LibreOffice 不可用 ⇒ 区域取数照常返回缓存值并带 warning「公式未能重算，显示的是文件缓存值」。
"""
from __future__ import annotations

import ast
import hashlib
import uuid
from collections.abc import AsyncIterator
from contextlib import asynccontextmanager
from pathlib import Path
from types import SimpleNamespace
from typing import Any

import pytest
from openpyxl import Workbook, load_workbook

import app.services.wp_template_init_service as wp_template_init_service
from app.models.workpaper_models import WorkingPaper, WpIndex, WpSourceType
from app.routers import custom_query
from app.services.custom_query.fetchers import workpaper_cells
from tests.authz_world import TARGET_PROJECT_ID as X
from tests.authz_world import pg_world

APP = Path(__file__).resolve().parents[2] / "app"
MODULE = APP / "services" / "custom_query" / "fetchers" / "workpaper_cells.py"
ROUTER = APP / "routers" / "custom_query.py"
STORAGE_HELPERS = frozenset({"init_workpaper_from_template", "get_workpaper_storage_path"})
SHUTIL_WRITERS = frozenset({"copy", "copy2", "copyfile", "copytree", "copyfileobj", "move", "rmtree"})
OS_WRITERS = frozenset({"replace", "rename", "renames", "remove", "unlink", "makedirs", "mkdir", "rmdir", "removedirs",
                        "symlink", "link", "truncate"})
METHOD_WRITERS = frozenset({"write_bytes", "write_text", "save", "unlink", "touch", "mkdir", "rmdir", "rename",
                            "symlink_to", "hardlink_to"})


# ─── 扫描器 ───────────────────────────────────────────────────────────────────


def _module_aliases(tree: ast.AST, module: str) -> tuple[set[str], set[str]]:
    """``import module [as x]`` 绑定的名字，与 ``from module import f [as g]`` 绑定的名字（任意作用域）。"""
    modules, members = set(), set()
    for node in ast.walk(tree):
        if isinstance(node, ast.Import):
            modules |= {alias.asname or alias.name for alias in node.names if alias.name == module}
        elif isinstance(node, ast.ImportFrom) and node.module == module:
            members |= {alias.asname or alias.name for alias in node.names}
    return modules, members


def _open_mode(call: ast.Call, index: int) -> ast.expr | None:
    for keyword in call.keywords:
        if keyword.arg == "mode":
            return keyword.value
    return call.args[index] if len(call.args) > index else None


def _writes_by_open(call: ast.Call) -> bool:
    func = call.func
    if isinstance(func, ast.Name) and func.id == "open":
        mode = _open_mode(call, 1)
    elif isinstance(func, ast.Attribute) and func.attr == "open":
        builtin_like = isinstance(func.value, ast.Name) and func.value.id in {"io", "codecs"}
        mode = _open_mode(call, 1 if builtin_like else 0)
    else:
        return False
    if mode is None:
        return False
    if not (isinstance(mode, ast.Constant) and isinstance(mode.value, str)):
        return True  # 模式不是字面量：无从判定，按写处理
    return any(flag in mode.value for flag in "wax+")


def write_calls(source: str) -> list[tuple[str, int, str]]:
    """源码里所有写盘调用：（所在函数名、行号、形态）。模块级的调用记作函数名 ``<module>``。"""
    tree = ast.parse(source)
    shutil_mods, shutil_members = _module_aliases(tree, "shutil")
    os_mods, os_members = _module_aliases(tree, "os")
    found: list[tuple[str, int, str]] = []

    def visit(node: ast.AST, owner: str) -> None:
        for child in ast.iter_child_nodes(node):
            name = child.name if isinstance(child, (ast.FunctionDef, ast.AsyncFunctionDef)) else owner
            if isinstance(child, ast.Call):
                func, form = child.func, None
                if isinstance(func, ast.Attribute):
                    base = func.value.id if isinstance(func.value, ast.Name) else None
                    if base in shutil_mods and func.attr in SHUTIL_WRITERS:
                        form = f"shutil.{func.attr}"
                    elif base in os_mods and func.attr in OS_WRITERS:
                        form = f"os.{func.attr}"
                    elif func.attr in METHOD_WRITERS:
                        form = f".{func.attr}"
                    elif func.attr == "replace" and len(child.args) == 1 and not child.keywords:
                        form = ".replace（Path.replace 只收一个实参；str.replace 至少两个）"
                elif isinstance(func, ast.Name):
                    if func.id in shutil_members & SHUTIL_WRITERS or func.id in os_members & OS_WRITERS:
                        form = func.id
                if form is None and _writes_by_open(child):
                    form = "open(写模式)"
                if form:
                    found.append((owner, child.lineno, form))
            visit(child, name)

    visit(tree, "<module>")
    return found


def _function(tree: ast.Module, name: str) -> ast.FunctionDef:
    matches = [n for n in ast.walk(tree) if isinstance(n, (ast.FunctionDef, ast.AsyncFunctionDef)) and n.name == name]
    assert len(matches) == 1, f"{MODULE.name} 里应恰有一个 {name}，实得 {len(matches)}"
    return matches[0]


def _tempdir_names(func: ast.AST) -> set[str]:
    """``with tempfile.TemporaryDirectory(...) as <名字>`` 在函数里绑定的名字。"""
    names = set()
    for node in ast.walk(func):
        if isinstance(node, (ast.With, ast.AsyncWith)):
            for item in node.items:
                call = item.context_expr
                if (isinstance(call, ast.Call) and isinstance(call.func, ast.Attribute)
                        and call.func.attr == "TemporaryDirectory" and isinstance(call.func.value, ast.Name)
                        and call.func.value.id == "tempfile" and isinstance(item.optional_vars, ast.Name)):
                    names.add(item.optional_vars.id)
    return names


def _is_tempdir_expr(expr: ast.expr, names: set[str]) -> bool:
    if isinstance(expr, ast.Name):
        return expr.id in names
    if isinstance(expr, ast.Call) and isinstance(expr.func, ast.Name) and expr.func.id in {"Path", "str"}:
        return len(expr.args) == 1 and _is_tempdir_expr(expr.args[0], names)
    return False


def _referenced_names(tree: ast.AST) -> set[str]:
    names = set()
    for node in ast.walk(tree):
        if isinstance(node, ast.Name):
            names.add(node.id)
        elif isinstance(node, ast.Attribute):
            names.add(node.attr)
        elif isinstance(node, (ast.Import, ast.ImportFrom)):
            names |= {alias.name.rsplit(".", 1)[-1] for alias in node.names}
        elif isinstance(node, ast.Constant) and isinstance(node.value, str):
            names.add(node.value)
    return names


# ─── 一、AST 守卫 ─────────────────────────────────────────────────────────────


def _module_tree() -> ast.Module:
    return ast.parse(MODULE.read_text(encoding="utf-8"))


def test_query_path_never_generates_files_from_templates() -> None:
    hits = sorted(_referenced_names(_module_tree()) & STORAGE_HELPERS)
    assert not hits, f"{MODULE.name} 不得引用会写 storage 的辅助（从模板生成文件 / mkdir）：{hits}"
    router_hits = sorted(_referenced_names(ast.parse(ROUTER.read_text(encoding="utf-8")))
                         & (STORAGE_HELPERS | {"get_workpaper_file"}))
    assert not router_hits, f"routers/custom_query.py 应经 fetchers/workpaper_cells 读底稿文件，不直接调：{router_hits}"


def test_disk_writes_only_inside_temp_copy() -> None:
    writes = write_calls(MODULE.read_text(encoding="utf-8"))
    outside = [w for w in writes if w[0] != "_temp_copy"]
    assert not outside, f"查询路径的写盘只许出现在 _temp_copy 里（R4.4）：{outside}"
    assert [w[2] for w in writes] == ["shutil.copy2"], f"_temp_copy 应恰有一次 shutil.copy2：{writes}"


def test_temp_copy_destination_is_always_a_tempfile_directory() -> None:
    tree = _module_tree()
    sites = []
    for func in (n for n in ast.walk(tree) if isinstance(n, (ast.FunctionDef, ast.AsyncFunctionDef))):
        names = _tempdir_names(func)
        for call in (n for n in ast.walk(func) if isinstance(n, ast.Call)):
            if isinstance(call.func, ast.Name) and call.func.id == "_temp_copy":
                target = call.args[1] if len(call.args) > 1 else next(
                    (k.value for k in call.keywords if k.arg == "tmpdir"), None)
                sites.append((func.name, call.lineno, target is not None and _is_tempdir_expr(target, names)))
    assert sites, "应至少有一处调用 _temp_copy（否则本守卫空转）"
    bad = [s for s in sites if not s[2]]
    assert not bad, f"_temp_copy 的目标目录必须是 tempfile.TemporaryDirectory() 绑定的目录：{bad}"


def test_recompute_returns_bytes_and_runs_libreoffice_on_the_temp_copy() -> None:
    func = _function(_module_tree(), "recompute_with_libreoffice")
    assert func.returns is not None and ast.unparse(func.returns) in {"bytes | None", "Optional[bytes]"}, (
        f"recompute_with_libreoffice 应标注返回 bytes | None，实得 {ast.unparse(func.returns) if func.returns else None}")
    for ret in (n for n in ast.walk(func) if isinstance(n, ast.Return)):
        value = ret.value
        ok = (value is None or (isinstance(value, ast.Constant) and value.value is None)
              or (isinstance(value, ast.Call) and isinstance(value.func, ast.Attribute)
                  and value.func.attr == "read_bytes"))
        assert ok, f"recompute_with_libreoffice 只许返回 None 或 ….read_bytes()：L{ret.lineno} {ast.unparse(ret)}"
    assert not (_referenced_names(func) & (STORAGE_HELPERS | {"get_workpaper_file", "existing_workpaper_file",
                                                              "STORAGE_DIR"})), "重算不得触碰 storage"
    copies = {t.id for n in ast.walk(func) if isinstance(n, ast.Assign) and isinstance(n.value, ast.Call)
              and isinstance(n.value.func, ast.Name) and n.value.func.id == "_temp_copy" for t in n.targets
              if isinstance(t, ast.Name)}
    names = _tempdir_names(func)
    runs = [n for n in ast.walk(func) if isinstance(n, ast.Call) and isinstance(n.func, ast.Attribute)
            and n.func.attr in {"run", "Popen", "call", "check_call", "check_output"}]
    assert len(runs) == 1 and runs[0].args and isinstance(runs[0].args[0], ast.List), "应恰有一次以列表传参的子进程调用"
    argv = runs[0].args[0].elts
    flags = [e.value if isinstance(e, ast.Constant) else None for e in argv]
    assert "--outdir" in flags, "LibreOffice 须显式 --outdir"
    outdir = argv[flags.index("--outdir") + 1]
    assert _is_tempdir_expr(outdir, names), f"--outdir 必须是临时目录：{ast.unparse(outdir)}"
    source = argv[-1]
    inner = source.args[0] if isinstance(source, ast.Call) and len(source.args) == 1 else source
    assert isinstance(inner, ast.Name) and inner.id in copies, (
        f"LibreOffice 的输入必须是 _temp_copy 的产物（不是原文件）：{ast.unparse(source)}")


def test_scanner_recognises_every_write_form_and_ignores_reads() -> None:
    """扫描器自检（双向）：写盘形态逐个认得，相近的只读写法不误认。"""
    positive = {
        "shutil.copy": "import shutil\ndef f(a, b):\n    shutil.copy(a, b)\n",
        "aliased shutil.move": "def f(a, b):\n    import shutil as _s\n    _s.move(a, b)\n",
        "from shutil import copy2": "from shutil import copy2\ndef f(a, b):\n    copy2(a, b)\n",
        "os.replace": "import os\ndef f(a, b):\n    os.replace(a, b)\n",
        "write_bytes": "def f(p):\n    p.write_bytes(b'')\n",
        "workbook save": "def f(wb, p):\n    wb.save(p)\n",
        "Path.replace": "def f(p, q):\n    p.replace(q)\n",
        "mkdir": "def f(p):\n    p.mkdir(parents=True)\n",
        "open wb": "def f(p):\n    open(p, 'wb')\n",
        "open mode kw": "def f(p):\n    open(p, mode='a')\n",
        "Path.open w": "def f(p):\n    p.open('w')\n",
        "open dynamic mode": "def f(p, m):\n    open(p, m)\n",
    }
    for label, source in positive.items():
        assert [w[0] for w in write_calls(source)] == ["f"], f"漏认写盘形态：{label}"
    negative = (
        "def f(s):\n    return s.replace('a', 'b')\n",
        "def f(p):\n    return open(p, 'rb').read()\n",
        "def f(p):\n    return p.read_bytes()\n",
        "import shutil\ndef f():\n    return shutil.which('soffice')\n",
        "def f(d):\n    return d.copy()\n",
    )
    for source in negative:
        assert write_calls(source) == [], f"误认：{source!r}"


# ─── 二、行为（真 PG 权限世界 + 临时存储根）─────────────────────────────────────


def _storage_state(root: Path) -> list[tuple[str, str | None, int]]:
    """storage 下每个条目：（相对路径、文件 sha256 / 目录为 None、mtime_ns）。"""
    if not root.exists():
        return []
    return sorted(
        (p.relative_to(root).as_posix(), hashlib.sha256(p.read_bytes()).hexdigest() if p.is_file() else None,
         p.stat().st_mtime_ns)
        for p in root.rglob("*")
    )


@asynccontextmanager
async def _paper(monkeypatch: pytest.MonkeyPatch, tmp_path: Path, *, with_file: bool) -> AsyncIterator[dict[str, Any]]:
    """权限世界里 X 有一张没有 univer_snapshot 的底稿 D2（取数会落到 storage xlsx）；存储根换成本例临时目录。"""
    async with pg_world(monkeypatch, extra_tables=(WpIndex.__table__, WorkingPaper.__table__)) as world:
        admin = (await world.admin_non_member()).user
        index = WpIndex(id=uuid.uuid4(), project_id=X, wp_code="D2", wp_name="【测试】只读查询底稿")
        world.db.add(index)
        await world.db.flush()
        paper = WorkingPaper(id=uuid.uuid4(), project_id=X, wp_index_id=index.id, source_type=WpSourceType.template,
                             file_path=f"storage/projects/{X}/workpapers/D2.xlsx", parsed_data={})
        world.db.add(paper)
        await world.db.flush()
        storage = tmp_path / "wp_storage"
        monkeypatch.setattr(wp_template_init_service, "STORAGE_DIR", storage)
        if with_file:
            path = storage / str(X) / f"{paper.id}.xlsx"
            path.parent.mkdir(parents=True)
            book = Workbook()
            sheet = book.active
            sheet.title = "审定表"
            sheet["A1"], sheet["B1"] = 21, "=A1*2"  # B1 只有公式、没有缓存值 ⇒ 区域取数要重算
            book.save(path)
        async with world.client(custom_query.router, user=admin) as client:
            yield {"client": client, "storage": storage, "paper": paper}


async def _cell_range(client: Any) -> Any:
    return await client.post("/api/custom-query/execute", json={
        "project_id": str(X), "year": 2025, "source": "workpaper:D2|审定表", "filters": {"cell_range": "A1:B1"},
        "limit": 100,
    })


def _cells(response: Any) -> dict[str, Any]:
    assert response.status_code == 200, response.text[:800]
    return {row["cell_ref"]: row for row in response.json()["rows"]}


pg = pytest.mark.pg_only


@pg
async def test_sheet_preview_reads_existing_file_without_touching_it(monkeypatch, tmp_path) -> None:
    async with _paper(monkeypatch, tmp_path, with_file=True) as ctx:
        before = _storage_state(ctx["storage"])
        response = await ctx["client"].get(f"/api/custom-query/wp-sheet-preview?project_id={X}&wp_code=D2")
        data = response.json()
        assert response.status_code == 200 and data["available"] is True and data["source"] == "xlsx_cache", data
        assert data["cells"]["0,0"]["v"] == 21 and data["has_unresolved_formula"] is True
        assert _storage_state(ctx["storage"]) == before, "预览不得改写 storage（字节 / mtime / 文件集合）"


@pg
async def test_sheet_preview_without_file_creates_nothing(monkeypatch, tmp_path) -> None:
    async with _paper(monkeypatch, tmp_path, with_file=False) as ctx:
        response = await ctx["client"].get(f"/api/custom-query/wp-sheet-preview?project_id={X}&wp_code=D2")
        assert response.status_code == 200
        assert response.json() == {"available": False, "reason": workpaper_cells.NO_FILE_REASON, "wp_code": "D2"}
        assert not ctx["storage"].exists(), f"无文件时不得在 storage 生成任何东西：{_storage_state(ctx['storage'])}"


@pg
async def test_cell_range_recompute_runs_on_a_temp_copy_only(monkeypatch, tmp_path) -> None:
    calls: list[list[str]] = []

    def fake_run(args: list[str], **_kwargs: Any) -> SimpleNamespace:
        """LibreOffice 替身：把重算后的值写进 --outdir 下的输出副本（与真 soffice --convert-to 同一落点）。"""
        calls.append([str(a) for a in args])
        source, outdir = Path(args[-1]), Path(args[args.index("--outdir") + 1])
        book = load_workbook(source)
        book["审定表"]["B1"] = 42
        book.save(outdir / f"{source.stem}.xlsx")
        return SimpleNamespace(returncode=0, stderr=b"")

    monkeypatch.setattr(workpaper_cells, "_find_soffice", lambda: "soffice-fake")
    monkeypatch.setattr(workpaper_cells, "subprocess", SimpleNamespace(run=fake_run))
    async with _paper(monkeypatch, tmp_path, with_file=True) as ctx:
        before = _storage_state(ctx["storage"])
        cells = _cells(await _cell_range(ctx["client"]))
        assert (cells["A1"]["value"], cells["B1"]["value"], cells["B1"]["formula"]) == (21, 42, "=A1*2"), cells
        assert len(calls) == 1, calls
        storage = str(ctx["storage"].resolve())
        source, outdir = calls[0][-1], calls[0][calls[0].index("--outdir") + 1]
        assert not str(Path(source).resolve()).startswith(storage), f"LibreOffice 的输入不得是 storage 里的原文件：{source}"
        assert not str(Path(outdir).resolve()).startswith(storage), f"LibreOffice 不得输出到 storage：{outdir}"
        assert _storage_state(ctx["storage"]) == before, "重算不得改写 storage 里的原文件（字节 / mtime / 文件集合）"


@pg
async def test_cell_range_without_libreoffice_keeps_cached_values_and_warns(monkeypatch, tmp_path) -> None:
    monkeypatch.setattr(workpaper_cells, "_find_soffice", lambda: None)
    async with _paper(monkeypatch, tmp_path, with_file=True) as ctx:
        before = _storage_state(ctx["storage"])
        response = await _cell_range(ctx["client"])
        cells = _cells(response)
        assert (cells["A1"]["value"], cells["B1"]["value"]) == (21, ""), cells
        assert workpaper_cells.RECOMPUTE_FAILED_WARNING in response.json()["warnings"], response.json()["warnings"]
        assert _storage_state(ctx["storage"]) == before


@pg
async def test_cell_range_without_file_is_400_and_creates_nothing(monkeypatch, tmp_path) -> None:
    async with _paper(monkeypatch, tmp_path, with_file=False) as ctx:
        response = await _cell_range(ctx["client"])
        assert response.status_code == 400, response.text[:800]
        assert response.json()["detail"] == {"error_code": "QUERY_FAILED", "message": workpaper_cells.NO_FILE_REASON,
                                             "source": "workpaper:D2|审定表"}
        assert not ctx["storage"].exists(), f"无文件时不得在 storage 生成任何东西：{_storage_state(ctx['storage'])}"

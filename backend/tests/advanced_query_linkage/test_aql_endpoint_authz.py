"""端点权限配方表（advanced-query-linkage-closure 任务 2；design §10.3；R12.5、N2）。

表驱动：design §10.3 的第 1~29 行、第 29b 行与循环受限子表（23c / 25c / 26c / 28c / 23n）逐格一条用例；第 30~32 行（W3）
在任务 41 补。每格一个独立的权限世界（``tests/authz_world.py``），经生产注册的 router 用 httpx + ASGI 真发请求，只 override
``get_current_user`` 与 ``get_db``，门禁全是生产对象（``GateClient`` 每个请求前自检）。

- 身份（每个世界都建好，按列选一个发请求）：匿名 = 不登录；A = 系统 auditor、无成员行；R = X 上（readonly, readonly）；
  E = X 上（auditor, edit）；PM = X 上（manager, edit）、系统 auditor，另一条人员记录 + X 上 manager 委派（链路唯一的 X 主管）；
  SM = 系统 manager、只在 Y 有（manager, edit）成员行，另一条人员记录 + Y 上 manager 委派（Y 主管）；P = 系统 partner；
  AD = 系统 admin、非成员；CR = X 上（manager, edit）+ auditor 委派（受限、范围为空，只测子表）；NM = 系统 manager、
  没有任何成员行与人员记录（只测 23n）。``eqcr_assignment`` 每个身份只调一次（两条人员记录 ⇒ 链路不唯一 ⇒ 受限）。
- 模板：T = AD 公开；S = PM 共享给 X（scope=project）；Q = AD 私有。
- 底座：S = ``sqlite_world``（纯门禁格）；P = ``pg_world``（依赖文本 SQL 定位或断言结果内容的格；SQLite 存 32 位十六进制 UUID，
  文本 SQL 绑定 ``str(uuid)`` 永远匹配不上会假绿）。P 格带 ``pg_only``，跑真库时 skipped 须为 0。所需业务表经 ``extra_tables``
  补建，缺表即该格失败、不降级跳过；PM / SM 走到循环范围的格（第 1、4、5、13、23~26、28 行）补 ``wp_index``。
- 回写两类请求一律带 ``X-File-Opened-At``（缺头即 400，会掩盖门禁结论）；取明天，乐观锁不会碰巧给出 409。
- design 表格只写状态码的格，凡结果内容能区分对错的都加内容断言：业务视图 / 模板执行结果非空且只含 X；构建器结果的项目
  集合恰为应见集合；403 区分「没有构建器准入」与「人员表角色门」；回写预览与确认核对条目原因与错误码；回收站三行的前置
  先确认是软删（否则不可见列的格子会因「硬删后本来就不存在」而碰巧通过）。

现已成立的格直接断言；待修复的格登记在 ``PENDING``（行号 → 列 → 修复任务号与原因），以 ``xfail(strict=True)`` 钉住：
修好即 XPASS 变红，逼着删掉登记。路由对账：custom_query、custom_query_writeback、query_builder 三个 router 的每条路由
都要在表里出现，缺一条即红并点名（新增路由在其任务里补行；已删路由只出现在第 29 行，暂缓删除的 cross-sheet-trace 在 29b）。
"""
from __future__ import annotations

import io
import uuid
from collections.abc import Callable
from contextlib import asynccontextmanager
from dataclasses import dataclass, field
from datetime import datetime, timedelta, timezone
from decimal import Decimal
from pathlib import Path
from typing import Any

import pytest
from fastapi import APIRouter
from fastapi.routing import APIRoute
from httpx import Response
from openpyxl import load_workbook
from sqlalchemy import Table, select

import app.services.wp_template_init_service as wp_template_init_service
from app.models.audit_platform_models import AccountCategory, TrialBalance
from app.services.custom_query.fetchers import workpaper_cells
from app.models.base import UserRole
from app.models.core import User
from app.models.custom_query_models import CustomQueryTemplate
from app.models.workpaper_models import WorkingPaper, WpIndex, WpSourceType
from app.routers import custom_query, custom_query_writeback, query_builder, recycle_bin
from tests.authz_world import OTHER_PROJECT_ID, TARGET_PROJECT_ID, AuthzWorld, pg_world, sqlite_world

X, Y = TARGET_PROJECT_ID, OTHER_PROJECT_ID
YEAR = 2025
COLUMNS = ("anon", "A", "R", "E", "PM", "SM", "P", "AD")
RECONCILED_ROUTERS: tuple[APIRouter, ...] = (custom_query.router, custom_query_writeback.router, query_builder.router)
OPENED_AT = {"X-File-Opened-At": (datetime.now(timezone.utc) + timedelta(days=1)).isoformat()}
CYCLE_REASON = "您可访问的项目均设有审计循环范围，请在业务视图中查询"
NO_PROJECT_REASON = "您当前没有可访问的项目"


# ─── 世界里的数据 ─────────────────────────────────────────────────────────────


@dataclass
class Ctx:
    """一格用例的世界：身份、模板与底稿 id、临时存储根。"""

    world: AuthzWorld
    users: dict[str, User | None] = field(default_factory=dict)
    ids: dict[str, str] = field(default_factory=dict)
    storage: Path | None = None
    live_projects: set[str] = field(default_factory=set)

    def fmt(self, text: str) -> str:
        return text.format(X=X, Y=Y, **self.ids)


async def _identities(world: AuthzWorld, extra: tuple[str, ...]) -> dict[str, User | None]:
    users: dict[str, User | None] = {"anon": None, "A": await world.user(UserRole.auditor)}
    users["R"] = (await world.readonly_member()).user
    users["E"] = (await world.member("edit", "auditor")).user
    pm = (await world.edit_member()).user
    await world.eqcr_assignment("manager", user=pm)
    users["PM"] = pm
    sm = await world.non_member_with_other_project(global_role=UserRole.manager)
    await world.eqcr_assignment("manager", user=sm.user, project=sm.other_project)
    users["SM"] = sm.user
    users["P"] = await world.user(UserRole.partner)
    users["AD"] = (await world.admin_non_member()).user
    if "CR" in extra:
        cr = (await world.member("edit", "manager")).user
        await world.eqcr_assignment("auditor", user=cr)
        users["CR"] = cr
    if "NM" in extra:
        users["NM"] = await world.user(UserRole.manager)
    return users


async def _add_workpaper(world: AuthzWorld, project_id: uuid.UUID, wp_code: str) -> WorkingPaper:
    index = WpIndex(id=uuid.uuid4(), project_id=project_id, wp_code=wp_code, wp_name=f"【测试】{wp_code} 权限世界底稿")
    world.db.add(index)
    await world.db.flush()
    paper = WorkingPaper(
        id=uuid.uuid4(), project_id=project_id, wp_index_id=index.id, source_type=WpSourceType.template,
        file_path=f"storage/projects/{project_id}/workpapers/{wp_code}.xlsx", parsed_data={},
    )
    world.db.add(paper)
    await world.db.flush()
    return paper


async def _add_tb_row(world: AuthzWorld, project_id: uuid.UUID, code: str) -> None:
    world.db.add(TrialBalance(
        project_id=project_id, year=YEAR, company_code="001", standard_account_code=code, account_name="库存现金",
        account_category=AccountCategory.asset, unadjusted_amount=Decimal("100.00"), audited_amount=Decimal("100.00"),
    ))
    await world.db.flush()


async def _add_template(world: AuthzWorld, owner: User, *, name: str, scope: str, shared: list[uuid.UUID]) -> str:
    template = CustomQueryTemplate(
        id=uuid.uuid4(), name=name, data_source="tb_detail", config={"source": "tb_detail", "year": YEAR},
        scope=scope, shared_project_ids=shared, creator_id=owner.id, created_by=owner.id,
    )
    world.db.add(template)
    await world.db.flush()
    return str(template.id)


async def _furnish(world: AuthzWorld, row: Row, request: pytest.FixtureRequest) -> Ctx:
    ctx = Ctx(world)
    ctx.users = await _identities(world, tuple(column for column in row.expect if column not in COLUMNS))
    if "wp_x" in row.furnish:
        ctx.ids["WP_X"] = str((await _add_workpaper(world, X, "D2")).id)
    if "wp_y" in row.furnish:
        ctx.ids["WP_Y"] = str((await _add_workpaper(world, Y, "E1")).id)
    if "tb_xy" in row.furnish:
        await _add_tb_row(world, X, "1001")
        await _add_tb_row(world, Y, "1002")
    if "templates" in row.furnish:
        users = ctx.users
        ctx.ids["T"] = await _add_template(world, users["AD"], name="【测试】T 公开", scope="public", shared=[])
        ctx.ids["S"] = await _add_template(world, users["PM"], name="【测试】S 共享给 X", scope="project", shared=[X])
        ctx.ids["Q"] = await _add_template(world, users["AD"], name="【测试】Q 私有", scope="private", shared=[])
    if "deleted_project" in row.furnish:
        await world.project(name="【测试】权限世界已删项目", is_deleted=True)
    if "storage" in row.furnish:
        # 底稿存储根换成本格的临时目录：只读预览若写文件（AQ-23），写在这里、可被断言，不落主树 storage
        ctx.storage = request.getfixturevalue("tmp_path") / "wp_storage"
        request.getfixturevalue("monkeypatch").setattr(wp_template_init_service, "STORAGE_DIR", ctx.storage)
    ctx.live_projects = {str(X), str(Y)}
    return ctx


async def _pm_deletes_s(ctx: Ctx) -> None:
    """第 14~16 行的前置：S 由作者 PM 经产品路径删除，且是软删（行仍在、is_deleted 为真）。

    只看回收站列表里「有没有 S」分不清「软删后按可见性滤掉」（对）与「硬删了根本不存在」（AQ-30 的缺陷）——
    后者会让不可见列的格子碰巧通过，所以先在这里确认软删。
    """
    async with ctx.world.client(custom_query.router, user=ctx.users["PM"]) as client:
        response = await client.delete(ctx.fmt("/api/custom-query/templates/{S}"))
    assert response.status_code == 204, f"前置：PM 删除 S 应 204，实得 {response.status_code}：{response.text[:300]}"
    stmt = select(CustomQueryTemplate).where(CustomQueryTemplate.id == uuid.UUID(ctx.ids["S"]))
    row = (await ctx.world.db.execute(stmt.execution_options(populate_existing=True))).scalar_one_or_none()
    assert row is not None, "前置：删除模板应为软删（进回收站），实际行已不存在（硬删除，AQ-30）"
    assert getattr(row, "is_deleted", None) is True, "前置：软删后 is_deleted 应为真"


BEFORE: dict[str, Callable[[Ctx], Any]] = {"pm_deletes_s": _pm_deletes_s}


# ─── 期望 ─────────────────────────────────────────────────────────────────────


@dataclass(frozen=True)
class Expect:
    status: int
    check: Callable[[Response, Ctx], None] | None = None


def _ids(values: Any) -> set[str]:
    return {str(value) for value in values}


def _sees(t: bool, s: bool, q: bool) -> Expect:
    def check(response: Response, ctx: Ctx) -> None:
        listed = _ids(item["id"] for item in response.json()["templates"])
        got = tuple(ctx.ids[key] in listed for key in ("T", "S", "Q"))
        assert got == (t, s, q), f"模板 T / S / Q 是否可见：期望 {(t, s, q)}，实得 {got}"
    return Expect(200, check)


def _recycle(contains: bool) -> Expect:
    def check(response: Response, ctx: Ctx) -> None:
        rows = [item for item in response.json()["items"] if item["id"] == ctx.ids["S"]]
        assert bool(rows) is contains, f"回收站里有没有已删的 S：期望 {contains}，实得 {bool(rows)}"
        if contains:
            assert (rows[0].get("can_restore"), rows[0].get("can_purge")) == (True, True), rows[0]
    return Expect(200, check)


def _x_trial_balance() -> Expect:
    """业务视图查 X 的试算表：非空，且只有 X 的那一行（科目 1001；Y 的那一行是 1002）——防「200 但空结果」的假绿。"""
    def check(response: Response, ctx: Ctx) -> None:
        rows = response.json()["rows"]
        values = [set(map(str, row.values())) for row in rows]
        assert rows and all("1001" in v for v in values) and not any("1002" in v for v in values), rows
    return Expect(200, check)


def _x_cycle_empty() -> Expect:
    """R / E（只有成员行、没有人员记录与委派 ⇒ 循环受限、范围为空，任务 23）：200 且没有行——与试算表页同口径（ADR-AQL-19）。"""
    def check(response: Response, ctx: Ctx) -> None:
        body = response.json()
        assert body["rows"] == [] and body["total"] == 0, body["rows"][:3]
    return Expect(200, check)


def _x_export_cycle_empty() -> Expect:
    """同 :func:`_x_cycle_empty` 的导出：文件只有「查询结果为空」一行。"""
    def check(response: Response, ctx: Ctx) -> None:
        sheet = load_workbook(io.BytesIO(response.content), read_only=True).worksheets[0]
        rows = [[c for c in r if c is not None] for r in sheet.iter_rows(values_only=True)][1:]
        assert rows == [["（查询结果为空）"]], rows
    return Expect(200, check)


_PROJECT_KEYS = {"X": X, "Y": Y}


def _rows_only(*keys: str, column: str = "project_id") -> Expect:
    """结果非空，且项目恰为 ``keys``（防「200 但空结果 / 多了别的项目」的假绿）。"""
    want = {str(_PROJECT_KEYS[key]) for key in keys}

    def check(response: Response, ctx: Ctx) -> None:
        got = _ids(row[column] for row in response.json()["rows"])
        assert got == want, f"结果应非空且恰含 {'、'.join(keys)}，实得 {sorted(got)}"
    return Expect(200, check)


def _x_export() -> Expect:
    """业务视图导出 X 的试算表（任务 26）：文件的数据行非空，且只有 X 的那一行（科目 1001）。"""
    def check(response: Response, ctx: Ctx) -> None:
        sheet = load_workbook(io.BytesIO(response.content), read_only=True).worksheets[0]
        rows = [list(r) for r in sheet.iter_rows(values_only=True)][1:]
        values = [set(map(str, r)) for r in rows]
        assert rows and all("1001" in v for v in values) and not any("1002" in v for v in values), rows
    return Expect(200, check)


def _export_only(*keys: str) -> Expect:
    want = {str(_PROJECT_KEYS[key]) for key in keys}

    def check(response: Response, ctx: Ctx) -> None:
        sheet = load_workbook(io.BytesIO(response.content), read_only=True).worksheets[0]
        rows = [list(r) for r in sheet.iter_rows(values_only=True)]
        index = next(i for i, name in enumerate(rows[0]) if str(name).endswith("project_id"))
        got = _ids(r[index] for r in rows[1:] if r and r[index] is not None)
        assert got == want, f"导出文件应非空且恰含 {'、'.join(keys)}，实得 {sorted(got)}"
    return Expect(200, check)


def _all_live_projects() -> Expect:
    def check(response: Response, ctx: Ctx) -> None:
        got = _ids(row["id"] for row in response.json()["rows"])
        assert got == ctx.live_projects, f"应为全部未删项目 {sorted(ctx.live_projects)}，实得 {sorted(got)}"
    return Expect(200, check)


def _empty_with_exclusion() -> Expect:
    def check(response: Response, ctx: Ctx) -> None:
        data = response.json()
        assert data["rows"] == [], data["rows"]
        assert any("已排除 1 个" in w for w in data.get("warnings") or []), data.get("warnings")
    return Expect(200, check)


def _access(allowed: bool, all_projects: bool, count: int | None, reason: str | None = None) -> Expect:
    def check(response: Response, ctx: Ctx) -> None:
        data = response.json()
        got = (data.get("allowed"), data.get("all_projects"), data.get("accessible_project_count"))
        assert got == (allowed, all_projects, count), f"access：期望 {(allowed, all_projects, count)}，实得 {got}"
        if reason is not None:
            assert data.get("reason") == reason, data.get("reason")
    return Expect(200, check)


def _preview_unavailable() -> Expect:
    def check(response: Response, ctx: Ctx) -> None:
        data = response.json()
        assert data.get("available") is False, data
        assert data.get("reason") == workpaper_cells.NO_FILE_REASON, data
        created = sorted(str(p) for p in ctx.storage.rglob("*")) if ctx.storage and ctx.storage.exists() else []
        # 只读预览不生成文件，也不建目录（查找文件不经会 mkdir 的 get_workpaper_storage_path）
        assert ctx.storage is not None and not ctx.storage.exists(), f"只读预览不应在 storage 留下任何东西：{created}"
    return Expect(200, check)


def _preview_blocked() -> Expect:
    def check(response: Response, ctx: Ctx) -> None:
        data = response.json()
        assert data["enter_confirmation"] is False, data
        assert [item["error"] for item in data["items"]] == ["target_not_in_project"], data["items"]
    return Expect(200, check)


def _error_code(status: int, code: str) -> Expect:
    def check(response: Response, ctx: Ctx) -> None:
        detail = response.json().get("detail")
        assert isinstance(detail, dict) and detail.get("error_code") == code, detail
    return Expect(status, check)


def codes(*cells: int | Expect) -> dict[str, Expect]:
    assert len(cells) == len(COLUMNS), cells
    return {column: cell if isinstance(cell, Expect) else Expect(cell) for column, cell in zip(COLUMNS, cells)}


# ─── 配方表 ───────────────────────────────────────────────────────────────────


@dataclass(frozen=True)
class Call:
    method: str
    route: str  # router 上的路径模板（路由对账）
    url: str  # 请求 URL，占位 {X} {Y} {T} {S} {Q} {WP_X} {WP_Y}
    body: Callable[[Ctx], Any] | None = None
    headers: dict[str, str] | None = None
    label: str = ""


@dataclass(frozen=True)
class Row:
    no: str
    base: str  # "S" = sqlite_world；"P" = pg_world
    calls: tuple[Call, ...]
    expect: dict[str, Expect]
    routers: tuple[APIRouter, ...]
    tables: tuple[Table, ...] = ()
    furnish: tuple[str, ...] = ()
    before: tuple[str, ...] = ()


CQ = "/api/custom-query"
QB = "/api/query"
_WP = (WpIndex.__table__, WorkingPaper.__table__)
_TB = (TrialBalance.__table__, WpIndex.__table__)
_TPL = (CustomQueryTemplate.__table__,)
_EXECUTE_X = {"project_id": str(X), "year": YEAR, "source": "tb_detail", "limit": 100}
_TB_DSL = {"table": "trial_balance", "fields": ["project_id", "standard_account_code"], "limit": 100}


def _body(value: Any) -> Callable[[Ctx], Any]:
    return lambda ctx: value


def _writeback_targets(ctx: Ctx) -> dict[str, Any]:
    return {"project_id": str(X), "targets": [{
        "module": "workpaper", "wp_id": ctx.ids["WP_Y"], "wp_code": "E1", "sheet_name": "审定表", "cell_ref": "B7",
        "new_value": 1,
    }]}


ROWS: tuple[Row, ...] = (
    Row("1", "S", (Call("GET", f"{CQ}/indicators", f"{CQ}/indicators?project_id={{X}}&depth=1"),),
        codes(401, 403, 200, 200, 200, 403, 200, 200), (custom_query.router,), (WpIndex.__table__,)),
    Row("2", "P", (Call("GET", f"{CQ}/wp-id-by-code", f"{CQ}/wp-id-by-code?project_id={{X}}&wp_code=D2"),),
        codes(401, 403, 200, 200, 200, 403, 200, 200), (custom_query.router,), _WP, ("wp_x",)),
    Row("3", "P", (Call("GET", f"{CQ}/wp-sheet-preview", f"{CQ}/wp-sheet-preview?project_id={{X}}&wp_code=D2"),),
        codes(401, 403, *([_preview_unavailable()] * 3), 403, *([_preview_unavailable()] * 2)),
        (custom_query.router,), _WP, ("wp_x", "storage")),
    # R / E 只有成员行 ⇒ 循环受限、范围为空 ⇒ 200 无行（任务 23）；PM 是 X 主管、P / AD 不限 ⇒ 只有 X 的那一行
    Row("4", "P", (Call("POST", f"{CQ}/execute", f"{CQ}/execute", _body(_EXECUTE_X)),),
        codes(401, 403, _x_cycle_empty(), _x_cycle_empty(), _x_trial_balance(), 403, *([_x_trial_balance()] * 2)),
        (custom_query.router,), _TB, ("tb_xy",)),
    Row("5", "P", (Call("POST", f"{CQ}/export", f"{CQ}/export", _body(_EXECUTE_X)),),
        codes(401, 403, _x_export_cycle_empty(), _x_export_cycle_empty(), _x_export(), 403, *([_x_export()] * 2)),
        (custom_query.router,), _TB, ("tb_xy",)),
    Row("6", "P", (Call("GET", f"{CQ}/templates", f"{CQ}/templates"),),
        codes(401, _sees(True, False, False), _sees(True, True, False), _sees(True, True, False),
              _sees(True, True, False), _sees(True, False, False), _sees(True, True, False), _sees(True, True, True)),
        (custom_query.router,), _TPL, ("templates",)),
    Row("7", "S", (Call("POST", f"{CQ}/templates", f"{CQ}/templates", _body({
        "name": "【测试】私有模板", "data_source": "tb_detail", "config": {"source": "tb_detail", "year": YEAR},
        "scope": "private"})),),
        codes(401, 201, 201, 201, 201, 201, 201, 201), (custom_query.router,), _TPL),
    Row("8", "S", (Call("POST", f"{CQ}/templates", f"{CQ}/templates", _body({
        "name": "【测试】项目共享模板", "data_source": "tb_detail", "config": {"source": "tb_detail", "year": YEAR},
        "scope": "project", "shared_project_ids": [str(X)]})),),
        codes(401, 403, 403, 201, 201, 403, 403, 201), (custom_query.router,), _TPL),
    Row("9", "P", (Call("GET", f"{CQ}/templates/{{template_id}}", f"{CQ}/templates/{{S}}"),),
        codes(401, 403, 200, 200, 200, 403, 200, 200), (custom_query.router,), _TPL, ("templates",)),
    Row("10", "P", (Call("GET", f"{CQ}/templates/{{template_id}}", f"{CQ}/templates/{{Q}}"),),
        codes(401, 403, 403, 403, 403, 403, 403, 200), (custom_query.router,), _TPL, ("templates",)),
    Row("11", "P", (Call("PUT", f"{CQ}/templates/{{template_id}}", f"{CQ}/templates/{{S}}",
                         _body({"name": "【测试】S 改名"})),),
        codes(401, 403, 403, 403, 200, 403, 403, 403), (custom_query.router,), _TPL, ("templates",)),
    Row("12", "P", (Call("DELETE", f"{CQ}/templates/{{template_id}}", f"{CQ}/templates/{{S}}"),),
        codes(401, 403, 403, 403, 204, 403, 403, 403), (custom_query.router,), _TPL, ("templates",)),
    # 年度随请求带上（界面执行时总有当前年度；不带时取目标项目审计年度，权限世界的项目没有审计年度 ⇒ 422，另有用例）
    Row("13", "P", (Call("POST", f"{CQ}/templates/{{template_id}}/execute", f"{CQ}/templates/{{S}}/execute",
                         _body({"project_id": str(X), "year": YEAR})),),
        codes(401, 403, _x_cycle_empty(), _x_cycle_empty(), _x_trial_balance(), 403, *([_x_trial_balance()] * 2)),
        (custom_query.router,), (*_TPL, *_TB), ("templates", "tb_xy")),
    Row("14", "P", (Call("GET", "/api/recycle-bin", "/api/recycle-bin?item_type=custom_query_template"),),
        codes(401, *([_recycle(False)] * 3), _recycle(True), _recycle(False), _recycle(False), _recycle(True)),
        (recycle_bin.router,), _TPL, ("templates",), ("pm_deletes_s",)),
    Row("15", "P", (Call("POST", "/api/recycle-bin/{item_type}/{item_id}/restore",
                         "/api/recycle-bin/custom_query_template/{S}/restore"),),
        codes(401, 403, 403, 403, 200, 403, 403, 200), (recycle_bin.router,), _TPL, ("templates",), ("pm_deletes_s",)),
    Row("16", "P", (Call("DELETE", "/api/recycle-bin/{item_type}/{item_id}", "/api/recycle-bin/custom_query_template/{S}"),),
        codes(401, 403, 403, 403, 200, 403, 403, 200), (recycle_bin.router,), _TPL, ("templates",), ("pm_deletes_s",)),
    Row("17", "S", (Call("POST", f"{CQ}/cell-writeback", f"{CQ}/cell-writeback", _body({
        "project_id": str(X), "wp_code": "D2", "sheet_name": "审定表", "cell_ref": "B7", "new_value": 1,
        "module": "tb"}), OPENED_AT),),
        codes(401, 422, 422, 422, 422, 422, 422, 422), (custom_query.router,)),
    Row("18", "P", (Call("POST", f"{CQ}/cell-writeback", f"{CQ}/cell-writeback", _body({
        "project_id": str(X), "wp_code": "E1", "sheet_name": "审定表", "cell_ref": "B7", "new_value": 1,
        "module": "workpaper"}), OPENED_AT),),
        codes(401, 403, 403, 404, 404, 403, 404, 404), (custom_query.router,), _WP, ("wp_y",)),
    Row("19", "P", (Call("POST", f"{CQ}/writeback-preview", f"{CQ}/writeback-preview", _writeback_targets),),
        codes(401, 403, 403, _preview_blocked(), _preview_blocked(), 403, _preview_blocked(), _preview_blocked()),
        (custom_query.router, custom_query_writeback.router), _WP, ("wp_y",)),
    Row("20", "P", (Call("POST", f"{CQ}/writeback-confirm", f"{CQ}/writeback-confirm", _writeback_targets, OPENED_AT),),
        codes(401, 403, 403, *([_error_code(409, "WRITEBACK_CONFLICT")] * 2), 403,
              *([_error_code(409, "WRITEBACK_CONFLICT")] * 2)),
        (custom_query.router, custom_query_writeback.router), _WP, ("wp_y",)),
    Row("21", "S", (Call("GET", f"{CQ}/index-health", f"{CQ}/index-health"),),
        codes(401, 200, 200, 200, 200, 200, 200, 200), (custom_query.router,)),
    Row("22", "S", (Call("GET", f"{CQ}/address-resolve", f"{CQ}/address-resolve?uri=workpaper:A1"),),
        codes(401, 200, 200, 200, 200, 200, 200, 200), (custom_query.router,)),
    Row("23", "S", (Call("GET", f"{QB}/access", f"{QB}/access"),),
        codes(401, _access(False, False, 0), _access(False, False, 0), _access(False, False, 0),
              _access(True, False, 1), _access(True, False, 1), _access(True, True, None), _access(True, True, None)),
        (query_builder.router,), (WpIndex.__table__,)),
    Row("24", "S", (Call("GET", f"{QB}/schema", f"{QB}/schema", label="schema"),
                    Call("POST", f"{QB}/preview", f"{QB}/preview", _body(_TB_DSL), label="preview")),
        codes(401, 403, 403, 403, 200, 200, 200, 200), (query_builder.router,), (WpIndex.__table__,)),
    Row("25", "P", (Call("POST", f"{QB}/execute", f"{QB}/execute", _body(_TB_DSL), label="execute"),),
        codes(401, 403, 403, 403, _rows_only("X"), _rows_only("Y"), _rows_only("X", "Y"), _rows_only("X", "Y")),
        (query_builder.router,), _TB, ("tb_xy",)),
    Row("25", "P", (Call("POST", f"{QB}/export-excel", f"{QB}/export-excel", _body(_TB_DSL), label="export"),),
        codes(401, 403, 403, 403, _export_only("X"), _export_only("Y"), _export_only("X", "Y"),
              _export_only("X", "Y")),
        (query_builder.router,), _TB, ("tb_xy",)),
    Row("26", "P", (Call("POST", f"{QB}/execute", f"{QB}/execute", _body(
        {"table": "projects", "fields": ["id", "name", "is_deleted"], "limit": 1000})),),
        codes(401, 403, 403, 403, _rows_only("X", column="id"), _rows_only("Y", column="id"), _all_live_projects(),
              _all_live_projects()),
        (query_builder.router,), _TB, ("deleted_project", "tb_xy")),
    # 403 分两种原因：没有构建器准入（ROLE_FORBIDDEN）与有准入但人员表只对 admin / partner 开放（TABLE_ROLE_FORBIDDEN）
    Row("27", "S", (Call("POST", f"{QB}/execute", f"{QB}/execute", _body(
        {"table": "staff_members", "fields": ["id", "name"], "limit": 100})),),
        codes(401, *([_error_code(403, "ROLE_FORBIDDEN")] * 3), *([_error_code(403, "TABLE_ROLE_FORBIDDEN")] * 2),
              200, 200),
        (query_builder.router,), (WpIndex.__table__,)),
    # 显式指定 Y：PM 不可访问 Y ⇒ 403 FORBIDDEN_PROJECT 且不执行；可访问 Y 的身份结果只含 Y（指定项目不能被忽略）
    Row("28", "P", (Call("POST", f"{QB}/execute", f"{QB}/execute", _body({**_TB_DSL, "project_id": str(Y)})),),
        codes(401, *([_error_code(403, "ROLE_FORBIDDEN")] * 3), _error_code(403, "FORBIDDEN_PROJECT"),
              *([_rows_only("Y")] * 3)),
        (query_builder.router,), _TB, ("tb_xy",)),
    Row("29", "S", (Call("POST", f"{CQ}/batch-execute", f"{CQ}/batch-execute", _body(
        {"wp_codes": ["D2"], "project_id": str(X), "year": YEAR})),),
        codes(404, 404, 404, 404, 404, 404, 404, 404), (custom_query.router,)),
    Row("29b", "S", (Call("GET", f"{CQ}/cross-sheet-trace",
                          f"{CQ}/cross-sheet-trace?project_id={{X}}&wp_code=D2&sheet_name=审定表&cell_ref=B7"),),
        codes(401, 403, 404, 404, 404, 403, 404, 404), (custom_query.router,), _WP),
    # 循环受限子表：CR（23c / 25c / 26c / 28c）对照 NM（23n，本来就没有可访问项目）
    Row("23c", "S", (Call("GET", f"{QB}/access", f"{QB}/access"),),
        {"CR": _access(True, False, 0, CYCLE_REASON)}, (query_builder.router,), (WpIndex.__table__,)),
    Row("25c", "P", (Call("POST", f"{QB}/execute", f"{QB}/execute", _body(_TB_DSL)),),
        {"CR": _empty_with_exclusion()}, (query_builder.router,), _TB, ("tb_xy",)),
    Row("26c", "P", (Call("POST", f"{QB}/execute", f"{QB}/execute", _body(
        {"table": "projects", "fields": ["id", "name", "is_deleted"], "limit": 1000})),),
        {"CR": _empty_with_exclusion()}, (query_builder.router,), (WpIndex.__table__,)),
    Row("28c", "P", (Call("POST", f"{QB}/execute", f"{QB}/execute", _body({**_TB_DSL, "project_id": str(X)})),),
        {"CR": _error_code(403, "CYCLE_SCOPED_PROJECT")}, (query_builder.router,), _TB, ("tb_xy",)),
    Row("23n", "S", (Call("GET", f"{QB}/access", f"{QB}/access"),),
        {"NM": _access(True, False, 0, NO_PROJECT_REASON)}, (query_builder.router,), (WpIndex.__table__,)),
)

#: 待修复的格：行号 → 列（"*" = 整行）→ 修复任务与原因（2026-10-27 B1 现跑的红形态）。修好一格删一格；全部转正在任务 35。
_LOGGED_IN = ("A", "R", "E", "PM", "SM", "P", "AD")


def _each(columns: tuple[str, ...], reason: str) -> dict[str, str]:
    return dict.fromkeys(columns, reason)


PENDING: dict[str, dict[str, str]] = {
    # 6~9、13 在任务 24（B8）转正
    # 14~16 的前置（PM 软删 S）对每一列都先执行，硬删时整行在前置处失败
    "14": {"*": "任务 25（AQ-30）：模板是硬删、回收站没有模板类型；另：合入前本地 HEAD 的回收站列表 import 缺失模块 "
                "app.services.user_display_names（远端已有，BM 合入后补齐）"},
    "15": {"*": "任务 25（AQ-30）：模板是硬删；回收站不支持 custom_query_template 类型（400）"},
    "16": {"*": "任务 25（AQ-30）：模板是硬删；回收站不支持 custom_query_template 类型（400）"},
    "22": _each(_LOGGED_IN, "任务 32（AQ-17）：wp_template_registry 表不存在时 address-resolve 500"),
    # 23~28、23n 在任务 9 / 12（B3）转正；循环受限子表 23c / 25c / 26c / 28c 在任务 23（B6）转正
}


def _pending(row: Row, column: str) -> str | None:
    reasons = PENDING.get(row.no, {})
    return reasons.get(column) or reasons.get("*")


def _cells() -> list[Any]:
    cells = []
    for row_index, row in enumerate(ROWS):
        for call_index, call in enumerate(row.calls):
            for column in row.expect:
                marks = [pytest.mark.pg_only] if row.base == "P" else []
                reason = _pending(row, column)
                if reason:
                    marks.append(pytest.mark.xfail(strict=True, reason=reason))
                suffix = f"·{call.label}" if call.label else ""
                cells.append(pytest.param(row_index, call_index, column, marks=marks, id=f"{row.no}{suffix}-{column}"))
    return cells


@asynccontextmanager
async def _world(row: Row, monkeypatch: pytest.MonkeyPatch):
    factory = pg_world if row.base == "P" else sqlite_world
    async with factory(monkeypatch, extra_tables=row.tables) as world:
        yield world


@pytest.mark.parametrize(("row_index", "call_index", "column"), _cells())
async def test_endpoint_authz_cell(
    row_index: int, call_index: int, column: str, monkeypatch: pytest.MonkeyPatch, request: pytest.FixtureRequest,
) -> None:
    row, call = ROWS[row_index], ROWS[row_index].calls[call_index]
    async with _world(row, monkeypatch) as world:
        ctx = await _furnish(world, row, request)
        for step in row.before:
            await BEFORE[step](ctx)
        async with world.client(*row.routers, user=ctx.users[column]) as client:
            response = await client.request(
                call.method, ctx.fmt(call.url), json=call.body(ctx) if call.body else None, headers=call.headers,
            )
        expect = row.expect[column]
        assert response.status_code == expect.status, (
            f"第 {row.no} 行 {call.method} {call.url} 身份 {column}：期望 {expect.status}，实得 {response.status_code}："
            f"{response.text[:500]}"
        )
        if expect.check is not None:
            expect.check(response, ctx)


# ─── 表本身 ───────────────────────────────────────────────────────────────────


def _declared_routes() -> set[tuple[str, str]]:
    return {(call.method, call.route) for row in ROWS for call in row.calls}


def _router_routes() -> set[tuple[str, str]]:
    found = set()
    for router in RECONCILED_ROUTERS:
        for route in router.routes:
            if isinstance(route, APIRoute):
                found |= {(method, route.path) for method in route.methods if method != "HEAD"}
    return found


def test_every_route_of_the_three_routers_is_in_the_table() -> None:
    missing = sorted(_router_routes() - _declared_routes())
    assert not missing, "以下路由不在端点权限配方表里（新增路由须在其任务里补行，design §10.3）：" + "；".join(
        f"{method} {path}" for method, path in missing)


def test_table_rows_match_design_shape() -> None:
    main = [row for row in ROWS if set(row.expect) == set(COLUMNS)]
    numbers = [row.no for row in main]
    expected = [str(n) for n in range(1, 30)] + ["29b"]
    assert sorted(set(numbers), key=lambda n: (int(n.rstrip("b")), n)) == expected, numbers
    sub = {row.no: set(row.expect) for row in ROWS if set(row.expect) != set(COLUMNS)}
    assert sub == {"23c": {"CR"}, "25c": {"CR"}, "26c": {"CR"}, "28c": {"CR"}, "23n": {"NM"}}, sub
    for no, reasons in PENDING.items():
        rows = [row for row in ROWS if row.no == no]
        assert rows, f"PENDING 登记了不存在的行 {no}"
        for column in reasons:
            assert column == "*" or any(column in row.expect for row in rows), f"PENDING 第 {no} 行登记了不存在的列 {column}"
        assert all(reason.startswith("任务 ") for reason in reasons.values()), f"PENDING 第 {no} 行须写明修复任务号"

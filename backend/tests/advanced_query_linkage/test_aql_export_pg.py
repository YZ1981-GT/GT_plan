"""服务端完整导出（真 PG；advanced-query-linkage-closure 任务 26；R8.1~R8.6；design §10.2「导出」行）。

金标准在测试内用独立 SQL 现算（``gold``，不调被测代码）；被测请求经生产注册的 router（``business`` / ``builder`` 夹具：真库
回滚事务会话、fakeredis）。写入只在回滚事务里、只写【UAT】子公司A；真实项目只读。

- 和平药房序时账按月（2025-01）导出：数据行数 = COUNT（openpyxl 只读读回）、表头 = 请求传入的中文、「查询说明」七项齐全且
  行数 = COUNT、RFC 5987 文件名；审计 ``custom_query.export`` 进了审计队列；临时文件已删；耗时记在 junit 属性（N1：10 万行 ≤ 60 s）。
- 全年导出（> 1,000,000 行）：400 ``EXPORT_ROW_HARD_LIMIT``（``actual_rows`` = COUNT），不开始取数、不产生文件，名额已释放。
- 分组 400 ``EXPORT_GROUPING_UNSUPPORTED``、未登记筛选键 400（与 execute 同一校验）：都在取名额之前。
- 回滚事务给【UAT】子公司A 造摘要以 ``= + - @`` 开头的序时账行：导出读回都是文本单元格，数值 −1 仍是数值（变异 #18）。
- 构建器导出同月：行数 = COUNT（不再受 1,000 行限制）、表头 = 传入的中文、sheet 名为表名。
- 构建器导出华氏调整分录：表名「调整分录（AJE/RJE）」含 '/'，sheet 名清洗为全角「／」、200、行数 = COUNT（B7 评审 B7-R1）。
评估批 5 的「和平药房序时账可完整导出 1,147,414 行」按 ADR-AQL-14 / R8.2 调整为「按月完整导出 + 全年超限 400」（W11）。
"""
from __future__ import annotations

import io
import json
import tempfile
import time
from datetime import date
from decimal import Decimal
from urllib.parse import quote

import pytest
from openpyxl import load_workbook

import app.core.redis as core_redis
from app.models.audit_platform_models import TbLedger
from app.services.audit_logger_enhanced import AUDIT_LOG_QUEUE_KEY
from app.services.custom_query import audit_helper, export_gate, paged_fetch
from tests.advanced_query_linkage.gold import AUDIT_YEAR, HP, HS, UA, active_sql

pytestmark = pytest.mark.pg_only

#: 夹具把审计写入换成空操作；审计用例换回真实实现（经 audit_logger 推进本例的 fakeredis 队列，不进共享 Redis）
_REAL_LOG_UNTHROTTLED = audit_helper._log_unthrottled

EXPORT = "/api/custom-query/export"
XLSX = "application/vnd.openxmlformats-officedocument.spreadsheetml.sheet"
JAN = ["2025-01-01", "2025-01-31"]
LEDGER_HEADERS = {
    "voucher_date": "凭证日期", "voucher_no": "凭证号", "account_code": "科目编码", "account_name": "科目名称",
    "debit_amount": "借方发生额", "credit_amount": "贷方发生额", "summary": "摘要",
}
INFO_KEYS = ["项目", "年度", "数据源", "条件", "行数", "导出时间", "导出人"]


@pytest.fixture(autouse=True)
def _gate_and_tempdir(monkeypatch, tmp_path):
    """临时文件写进本例私有目录（结束时核对已删）；每例前后导出闸都必须是空的。"""
    xlsx_tmp = tmp_path / "xlsx_tmp"
    xlsx_tmp.mkdir()
    monkeypatch.setattr(tempfile, "tempdir", str(xlsx_tmp))
    assert export_gate.in_flight_count() == 0, "上一个用例泄漏了导出名额"
    yield xlsx_tmp
    leaked = export_gate.in_flight_count()
    export_gate._IN_FLIGHT.clear()
    assert leaked == 0, f"本例结束时仍有 {leaked} 个导出名额未释放"
    assert not sorted(p.name for p in xlsx_tmp.glob("aqm_export_*")), "导出临时文件没有删除"


async def _ledger_count(gold, pid, *, month: bool) -> int:
    sql = f"SELECT count(*) FROM tb_ledger t WHERE {active_sql()}"
    if month:
        sql += " AND t.voucher_date BETWEEN DATE '2025-01-01' AND DATE '2025-01-31'"
    return int(await gold.scalar(sql, pid=pid, year=AUDIT_YEAR))


def _read(content: bytes) -> tuple[list[str], list, int, dict]:
    """（sheet 名、表头、数据行数、查询说明）——只读模式逐行计数，不把 10 万行装进内存。"""
    wb = load_workbook(io.BytesIO(content), read_only=True)
    try:
        rows = wb.worksheets[0].iter_rows(values_only=True)
        header = list(next(rows))
        count = sum(1 for _ in rows)
        notes = {}
        if len(wb.worksheets) > 1:
            notes = {str(r[0]): r[1] for r in wb.worksheets[1].iter_rows(values_only=True) if r and r[0] is not None}
        return wb.sheetnames, header, count, notes
    finally:
        wb.close()


def _month(**extra) -> dict:
    return {"project_id": str(HP), "year": AUDIT_YEAR, "source": "ledger_entries",
            "conditions": [{"field": "voucher_date", "op": "between", "value": JAN}], "headers": LEDGER_HEADERS, **extra}


async def test_hp_month_export_rows_headers_info_and_audit(business, gold, monkeypatch, record_property):
    monkeypatch.setattr(audit_helper, "_log_unthrottled", _REAL_LOG_UNTHROTTLED)
    expected = await _ledger_count(gold, HP, month=True)
    assert expected > 1000, "金标准应超过旧的 1,000 行导出上限"
    started = time.perf_counter()
    response = await business.http.post(EXPORT, json=_month(file_name="和平药房_序时账_2025-01"))
    seconds = time.perf_counter() - started
    record_property("export_rows", expected)
    record_property("export_seconds", round(seconds, 1))
    assert response.status_code == 200, response.text[:800]
    assert response.headers["content-type"].startswith(XLSX)
    assert "filename*=UTF-8''" + quote("和平药房_序时账_2025-01.xlsx", safe="") in response.headers["content-disposition"]
    sheets, header, count, notes = _read(response.content)
    assert count == expected, f"导出数据行 {count} ≠ 金标准 COUNT {expected}"
    assert header == list(LEDGER_HEADERS.values())
    assert sheets == ["高级查询结果", "查询说明"]
    assert list(notes)[:7] == INFO_KEYS
    assert notes["行数"] == expected and notes["年度"] == AUDIT_YEAR
    assert notes["数据源"] == "序时账明细（ledger_entries）"
    assert notes["条件"] == "凭证日期 介于 2025-01-01 至 2025-01-31"
    assert str(HP) in notes["项目"] and notes["导出人"] == "aql-b5-admin"
    queued = [json.loads(raw) for raw in await core_redis.redis_client.lrange(AUDIT_LOG_QUEUE_KEY, 0, -1)]
    exports = [e for e in queued if e["action_type"] == "custom_query.export"]
    assert len(exports) == 1, queued
    assert exports[0]["project_id"] == str(HP)
    assert (exports[0]["payload"]["source"], exports[0]["payload"]["row_count"]) == ("ledger_entries", expected)


async def test_hp_full_year_export_is_rejected_before_any_row_is_read(business, gold, monkeypatch):
    total = await _ledger_count(gold, HP, month=False)
    assert total > 1_000_000, "金标准应超过导出硬上限（全年口径变了就换一个超限的项目年度）"
    streamed: list[tuple] = []

    def _no_stream(*args, **kwargs):  # 取数的下一层：超限时根本不该走到这里（走到即中止，不真去读 100 万行）
        streamed.append(args)
        raise AssertionError("超过导出上限时不应开始取数")

    monkeypatch.setattr(paged_fetch, "stream_statement", _no_stream)
    response = await business.http.post(EXPORT, json={k: v for k, v in _month().items() if k != "conditions"})
    assert streamed == [], "超过导出上限时开始了取数（COUNT 预检未在取数之前拦截）"
    assert response.status_code == 400, response.text[:500]
    assert response.headers["content-type"].startswith("application/json")
    detail = response.json()["detail"]
    assert detail["error_code"] == "EXPORT_ROW_HARD_LIMIT"
    assert (detail["actual_rows"], detail["row_hard_limit"]) == (total, 1_000_000)


@pytest.mark.parametrize(("extra", "code"), [
    ({"group": {"dimensions": ["account_code"], "aggregates": []}}, "EXPORT_GROUPING_UNSUPPORTED"),
    ({"filters": {"debit_amount_min": 1}}, "UNKNOWN_FILTER_KEY"),
])
async def test_rejections_happen_before_taking_a_slot(business, monkeypatch, extra: dict, code: str):
    taken: list[str] = []
    real_acquire = export_gate.acquire_export_slot

    def _spy(user_id):
        taken.append(str(user_id))
        return real_acquire(user_id)

    monkeypatch.setattr("app.services.custom_query.query_export.acquire_export_slot", _spy)
    response = await business.http.post(EXPORT, json=_month(**extra))
    assert response.status_code == 400, response.text[:500]
    assert response.json()["detail"]["error_code"] == code
    assert taken == [], "参数类 400 应发生在取导出名额之前"


async def test_formula_like_text_is_written_as_text(business, pg_rollback_session):
    texts = ["=1+1", "+1", "-1", "@x"]
    for i, text in enumerate(texts, 1):
        pg_rollback_session.add(TbLedger(
            project_id=UA, year=AUDIT_YEAR, company_code="001", voucher_date=date(2025, 1, i), voucher_no=f"AQLB7-{i}",
            account_code="1001", account_name="库存现金", entry_seq=i, debit_amount=Decimal("-1.00"),
            credit_amount=Decimal("0.00"), summary=text,
        ))
    await pg_rollback_session.flush()
    response = await business.http.post(EXPORT, json={
        "project_id": str(UA), "year": AUDIT_YEAR, "source": "ledger_entries",
        "columns": ["voucher_no", "debit_amount", "summary"], "sort": [{"field": "voucher_no", "direction": "asc"}],
        "headers": {"voucher_no": "凭证号", "debit_amount": "=借方", "summary": "摘要"},
    })
    assert response.status_code == 200, response.text[:800]
    ws = load_workbook(io.BytesIO(response.content)).worksheets[0]  # 非只读：要看单元格类型
    grid = [[(c.value, c.data_type) for c in row] for row in ws.iter_rows()]
    assert grid[0] == [("凭证号", "s"), ("=借方", "s"), ("摘要", "s")]
    assert [row[2] for row in grid[1:]] == [(text, "s") for text in texts]
    assert [row[1] for row in grid[1:]] == [(-1, "n")] * len(texts)
    assert [row[0][0] for row in grid[1:]] == [f"AQLB7-{i}" for i in range(1, len(texts) + 1)]


async def test_builder_month_export_rows_equal_count(builder, gold):
    expected = await _ledger_count(gold, HP, month=True)
    fields = ["voucher_date", "voucher_no", "account_code", "debit_amount", "credit_amount", "summary"]
    headers = {f: LEDGER_HEADERS[f] for f in fields}
    response = await builder.request(builder.admin, "POST", "/export-excel", {
        "table": "tb_ledger", "fields": fields, "project_id": str(HP), "year": AUDIT_YEAR, "limit": 100,
        "filters": [{"field": "voucher_date", "op": "between", "value": JAN}], "headers": headers,
    })
    assert response.status_code == 200, response.text[:800]
    sheets, header, count, notes = _read(response.content)
    assert count == expected and count > 1000, (
        f"构建器导出 {count} 行 ≠ 金标准 {expected}（旧实现复用分页语句，只导出 limit 行、至多 1,000）")
    assert header == [headers[f] for f in fields]
    assert sheets == ["序时账（总账明细）", "查询说明"]
    assert notes["行数"] == expected and str(HP) in notes["项目"]


async def test_builder_adjustments_export_has_a_valid_sheet_title(builder, gold):
    """改前 openpyxl ``create_sheet("调整分录（AJE/RJE）")`` 抛 ValueError，构建器导出调整分录恒 500（B7 评审 B7-R1）；
    枚举列写枚举值（金标准 ``::text``），不是「AdjustmentType.aje」（B7 复测发现）。"""
    fields = ["adjustment_no", "adjustment_type", "review_status", "account_code"]
    expected = sorted(tuple(r) for r in await gold.rows(
        "SELECT adjustment_no, adjustment_type::text, review_status::text, account_code FROM adjustments "
        "WHERE project_id = :pid AND year = :year AND NOT is_deleted", pid=HS, year=AUDIT_YEAR))
    assert expected, "前置：华氏 2025 有调整分录（没有就换一个有调整分录的项目年度）"
    response = await builder.request(builder.admin, "POST", "/export-excel", {
        "table": "adjustments", "fields": fields, "project_id": str(HS), "year": AUDIT_YEAR, "limit": 100})
    assert response.status_code == 200, response.text[:800]
    sheets, header, count, notes = _read(response.content)
    assert sheets == ["调整分录（AJE／RJE）", "查询说明"]
    assert header == fields and count == len(expected) and notes["行数"] == len(expected) and str(HS) in notes["项目"]
    wb = load_workbook(io.BytesIO(response.content), read_only=True)
    try:
        exported = sorted(tuple(r) for r in list(wb.worksheets[0].iter_rows(values_only=True))[1:])
    finally:
        wb.close()
    assert exported == expected

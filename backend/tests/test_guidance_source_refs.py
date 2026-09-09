"""source_ref resolver 的真实文件与 fail-closed 行为契约。"""
from __future__ import annotations

from pathlib import Path

from docx import Document
from docx.oxml import OxmlElement
from docx.oxml.ns import qn
from openpyxl import Workbook

from app.services.guidance_source_refs import (
    SourceRefContext,
    TemplateAuthority,
    file_sha256,
    validate_source_ref,
    validate_source_refs,
)


def _xlsx(path: Path, *, sheet: str = "说明", value: str = "审计程序来源") -> None:
    path.parent.mkdir(parents=True, exist_ok=True)
    workbook = Workbook()
    worksheet = workbook.active
    worksheet.title = sheet
    worksheet["A1"] = value
    worksheet["B2"] = "=1+1"
    workbook.save(path)
    workbook.close()


def _docx(
    path: Path,
    *,
    duplicate_bookmark: bool = False,
    close_bookmark: bool = True,
) -> None:
    path.parent.mkdir(parents=True, exist_ok=True)
    document = Document()
    document.add_paragraph("编制说明前言")

    paragraph = document.add_paragraph()
    start = OxmlElement("w:bookmarkStart")
    start.set(qn("w:id"), "17")
    start.set(qn("w:name"), "source_anchor")
    paragraph._p.append(start)
    paragraph.add_run("唯一书签正文")
    if close_bookmark:
        end = OxmlElement("w:bookmarkEnd")
        end.set(qn("w:id"), "17")
        paragraph._p.append(end)

    if duplicate_bookmark:
        duplicate = document.add_paragraph()
        duplicate_start = OxmlElement("w:bookmarkStart")
        duplicate_start.set(qn("w:id"), "18")
        duplicate_start.set(qn("w:name"), "source_anchor")
        duplicate._p.append(duplicate_start)
        duplicate.add_run("重复书签正文")
        duplicate_end = OxmlElement("w:bookmarkEnd")
        duplicate_end.set(qn("w:id"), "18")
        duplicate._p.append(duplicate_end)

    document.add_paragraph("唯一段落锚点")
    table = document.add_table(rows=1, cols=1)
    table.cell(0, 0).text = "表格正文"
    document.save(path)


def _context(repo: Path, source: Path, *, wp_code: str = "D0") -> SourceRefContext:
    relative = source.relative_to(repo).as_posix()
    return SourceRefContext(
        target_wp_code=wp_code,
        template_authorities=(
            TemplateAuthority(canonical_path=relative, wp_codes=(wp_code,)),
        ),
        repo_root=repo,
        template_root=repo / "backend" / "wp_templates",
    )


def _ref(source: Path, repo: Path, *, kind: str, **locator) -> dict:
    return {
        "kind": kind,
        "path": source.relative_to(repo).as_posix(),
        "digest": file_sha256(source),
        **locator,
    }


def _issue(result) -> str:
    assert result.issues
    return result.issues[0].code


def test_xlsx_validates_repo_path_digest_sheet_range_and_content(tmp_path: Path):
    source = tmp_path / "backend" / "wp_templates" / "D" / "D0.xlsx"
    _xlsx(source)
    context = _context(tmp_path, source)

    result = validate_source_ref(
        _ref(source, tmp_path, kind="xlsx", sheet="说明", range="A1:B2"),
        context,
    )

    assert result.status == "valid"
    assert result.facts.normalized_path == "backend/wp_templates/D/D0.xlsx"
    assert result.facts.observed_digest == file_sha256(source)
    assert result.facts.sheet == "说明"
    assert result.facts.cell_range == "A1:B2"
    assert len(result.facts.located_content_digest or "") == 64


def test_xlsx_rejects_path_escape_cross_template_and_digest_drift(tmp_path: Path):
    source = tmp_path / "backend" / "wp_templates" / "D" / "D0.xlsx"
    other = tmp_path / "backend" / "wp_templates" / "E" / "E0.xlsx"
    _xlsx(source)
    _xlsx(other)
    context = _context(tmp_path, source)

    escaped = validate_source_ref(
        {
            "kind": "xlsx",
            "path": "backend/wp_templates/../../outside.xlsx",
            "digest": "0" * 64,
            "sheet": "说明",
            "range": "A1",
        },
        context,
    )
    cross = validate_source_ref(
        _ref(other, tmp_path, kind="xlsx", sheet="说明", range="A1"),
        context,
    )
    stale = validate_source_ref(
        {
            **_ref(source, tmp_path, kind="xlsx", sheet="说明", range="A1"),
            "digest": "0" * 64,
        },
        context,
    )

    assert (escaped.status, _issue(escaped)) == ("invalid", "path_boundary_rejected")
    assert (cross.status, _issue(cross)) == (
        "cross_template",
        "template_not_active_for_context",
    )
    assert (stale.status, _issue(stale)) == ("stale", "digest_mismatch")


def test_xlsx_locator_is_required_and_must_still_resolve(tmp_path: Path):
    source = tmp_path / "backend" / "wp_templates" / "D" / "D0.xlsx"
    _xlsx(source)
    context = _context(tmp_path, source)
    base = _ref(source, tmp_path, kind="xlsx")

    no_sheet = validate_source_ref({**base, "range": "A1"}, context)
    no_range = validate_source_ref({**base, "sheet": "说明"}, context)
    wrong_sheet = validate_source_ref(
        {**base, "sheet": "已删除页", "range": "A1"}, context
    )
    empty_range = validate_source_ref(
        {**base, "sheet": "说明", "range": "C3:D4"}, context
    )
    oversized = validate_source_ref(
        {**base, "sheet": "说明", "range": "A1:XFD1048576"}, context
    )

    assert (no_sheet.status, _issue(no_sheet)) == ("invalid", "sheet_missing")
    assert (no_range.status, _issue(no_range)) == ("invalid", "range_missing")
    assert (wrong_sheet.status, _issue(wrong_sheet)) == ("stale", "sheet_not_found")
    assert (empty_range.status, _issue(empty_range)) == ("stale", "range_empty")
    assert (oversized.status, _issue(oversized)) == ("invalid", "range_too_large")


def test_docx_supports_bookmark_paragraph_and_before_first_table(tmp_path: Path):
    source = tmp_path / "backend" / "wp_templates" / "D" / "D0.docx"
    _docx(source)
    context = _context(tmp_path, source)
    base = _ref(source, tmp_path, kind="docx")

    bookmark = validate_source_ref(
        {**base, "anchor": {"type": "bookmark", "name": "source_anchor"}},
        context,
    )
    paragraph = validate_source_ref(
        {
            **base,
            "anchor": {
                "type": "paragraph_text",
                "text": "唯一段落锚点",
                "match": "exact",
            },
        },
        context,
    )
    before_table = validate_source_ref(
        {**base, "anchor": {"type": "before_first_table"}}, context
    )

    assert bookmark.status == "valid"
    assert bookmark.facts.anchor_match_count == 1
    assert bookmark.facts.anchor_type == "bookmark"
    assert paragraph.status == "valid"
    assert paragraph.facts.anchor_type == "paragraph_text"
    assert before_table.status == "valid"
    assert before_table.facts.anchor_type == "before_first_table"
    assert len(before_table.facts.located_content_digest or "") == 64


def test_docx_anchor_missing_duplicate_unclosed_and_local_digest_are_observable(tmp_path: Path):
    good = tmp_path / "backend" / "wp_templates" / "D" / "good.docx"
    duplicate = tmp_path / "backend" / "wp_templates" / "D" / "duplicate.docx"
    unclosed = tmp_path / "backend" / "wp_templates" / "D" / "unclosed.docx"
    _docx(good)
    _docx(duplicate, duplicate_bookmark=True)
    _docx(unclosed, close_bookmark=False)

    good_context = _context(tmp_path, good)
    missing = validate_source_ref(
        {
            **_ref(good, tmp_path, kind="docx"),
            "anchor": {"type": "bookmark", "name": "不存在"},
        },
        good_context,
    )
    duplicate_result = validate_source_ref(
        {
            **_ref(duplicate, tmp_path, kind="docx"),
            "anchor": {"type": "bookmark", "name": "source_anchor"},
        },
        _context(tmp_path, duplicate),
    )
    unclosed_result = validate_source_ref(
        {
            **_ref(unclosed, tmp_path, kind="docx"),
            "anchor": {"type": "bookmark", "name": "source_anchor"},
        },
        _context(tmp_path, unclosed),
    )
    local_stale = validate_source_ref(
        {
            **_ref(good, tmp_path, kind="docx"),
            "anchor": {
                "type": "bookmark",
                "name": "source_anchor",
                "content_digest": "0" * 64,
            },
        },
        good_context,
    )

    assert (missing.status, _issue(missing)) == ("stale", "anchor_not_found")
    assert (duplicate_result.status, _issue(duplicate_result)) == (
        "invalid",
        "anchor_ambiguous",
    )
    assert (unclosed_result.status, _issue(unclosed_result)) == (
        "invalid",
        "bookmark_unclosed",
    )
    assert (local_stale.status, _issue(local_stale)) == (
        "stale",
        "anchor_content_digest_mismatch",
    )


def test_docx_does_not_accept_legacy_range_as_semantic_anchor(tmp_path: Path):
    source = tmp_path / "backend" / "wp_templates" / "D" / "D0.docx"
    _docx(source)
    result = validate_source_ref(
        {
            **_ref(source, tmp_path, kind="docx"),
            "range": "before_first_table",
        },
        _context(tmp_path, source),
    )

    assert (result.status, _issue(result)) == ("invalid", "anchor_missing")


def test_authority_digest_and_active_override_are_part_of_validation(tmp_path: Path):
    canonical = tmp_path / "backend" / "wp_templates" / "D" / "D0.xlsx"
    active = tmp_path / "backend" / "storage" / "projects" / "p1" / "D0.xlsx"
    _xlsx(canonical, value="canonical")
    _xlsx(active, value="active override")
    authority = TemplateAuthority(
        canonical_path=canonical.relative_to(tmp_path).as_posix(),
        active_path=active,
        wp_codes=("D0",),
        origin="project_override",
        expected_active_digest=file_sha256(active),
    )
    context = SourceRefContext(
        target_wp_code="D0",
        template_authorities=(authority,),
        repo_root=tmp_path,
        template_root=tmp_path / "backend" / "wp_templates",
    )
    valid = validate_source_ref(
        {
            "kind": "xlsx",
            "path": canonical.relative_to(tmp_path).as_posix(),
            "digest": file_sha256(active),
            "sheet": "说明",
            "range": "A1",
        },
        context,
    )
    changed_authority = SourceRefContext(
        target_wp_code="D0",
        template_authorities=(
            TemplateAuthority(
                **{
                    **authority.__dict__,
                    "expected_active_digest": "0" * 64,
                }
            ),
        ),
        repo_root=tmp_path,
        template_root=tmp_path / "backend" / "wp_templates",
    )
    stale = validate_source_ref(
        {
            "kind": "xlsx",
            "path": canonical.relative_to(tmp_path).as_posix(),
            "digest": file_sha256(active),
            "sheet": "说明",
            "range": "A1",
        },
        changed_authority,
    )

    assert valid.status == "valid"
    assert valid.facts.active_template_origin == "project_override"
    assert valid.facts.observed_digest == file_sha256(active)
    assert (stale.status, _issue(stale)) == ("stale", "authority_digest_mismatch")


def test_batch_empty_invalid_and_digest_are_deterministic(tmp_path: Path):
    source = tmp_path / "backend" / "wp_templates" / "D" / "D0.xlsx"
    _xlsx(source)
    context = _context(tmp_path, source)

    empty = validate_source_refs([], context)
    invalid = validate_source_refs("not-a-list", context)
    first = validate_source_refs(
        [_ref(source, tmp_path, kind="xlsx", sheet="说明", range="A1")],
        context,
    )
    second = validate_source_refs(
        [_ref(source, tmp_path, kind="xlsx", sheet="说明", range="A1")],
        context,
    )

    assert empty.status_counts == {"missing": 1}
    assert _issue(empty.results[0]) == "ref_list_empty"
    assert invalid.status_counts == {"invalid": 1}
    assert _issue(invalid.results[0]) == "ref_list_invalid"
    assert first.all_valid is True
    assert first.facts_digest == second.facts_digest

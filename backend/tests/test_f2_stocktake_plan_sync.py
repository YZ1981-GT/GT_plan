"""F2-22 监盘计划 docx ↔ 结构化字段同步."""

from __future__ import annotations

from pathlib import Path

from app.services.f2_stocktake_plan_sync import (
    create_g2_6_2_template_docx,
    extract_fields_from_docx,
    fill_plan_docx,
    merge_extracted_into_existing,
)


def test_f2_22_template_has_placeholders(tmp_path: Path) -> None:
    tpl = create_g2_6_2_template_docx(tmp_path / "t.docx")
    from docx import Document

    text = "\n".join(p.text for p in Document(str(tpl)).paragraphs)
    assert "${purpose}" in text
    assert "${inventoryComposition}" in text
    assert "G2-6-2" in text


def test_f2_22_fill_extract_roundtrip(tmp_path: Path) -> None:
    tpl = create_g2_6_2_template_docx(tmp_path / "tpl.docx")
    target = tmp_path / "filled.docx"
    sample = {
        "entityName": "测试股份有限公司",
        "auditYear": "2025",
        "bsDate": "2025-12-31",
        "purpose": "确定存货存在并为公司所有。",
        "scope": "公司本部存货；子公司不纳入本次监盘。",
        "warehouses": "一分厂、工业园",
        "countDate": "2025年12月31日至2026年1月1日",
        "auditors": "张三、李四",
        "clientStaff": "仓管王五",
        "assignment": "张三：一分厂；李四：工业园",
        "prep": "（1）索取盘点计划；（2）人员就位。",
        "inventoryComposition": "原材料占57%，产成品占39%。",
        "countMethod": "公司主盘，我所监盘并抽盘。",
        "requirements": "数量覆盖不少于50%，金额不少于70%。",
        "remoteWarehouse": "无异地代管",
        "fraudRisk": "未见重大舞弊风险",
        "expertNeeded": "不需要专家",
        "planDate": "2025-12-29",
        "teamName": "测试审计小组",
    }
    fill_plan_docx(tpl, target, sample)
    extracted = extract_fields_from_docx(target)
    for key, value in sample.items():
        assert extracted.get(key) == value, f"{key}: {extracted.get(key)!r} != {value!r}"


def test_merge_extracted_keeps_unrelated_keys() -> None:
    existing = {"purpose": "旧目的", "customExtra": "保留"}
    extracted = {"purpose": "新目的", "scope": "新范围"}
    merged = merge_extracted_into_existing(existing, extracted)
    assert merged["purpose"] == "新目的"
    assert merged["scope"] == "新范围"
    assert merged["customExtra"] == "保留"

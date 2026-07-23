"""B2 前任沟通信函中文【】占位符自动填充测试。

Validates:
- P4 中文占位符替换正确（【被审计单位名称】/【20××】/【前任会计师事务所的名称】）
- P2 空值保留（无 B2-predecessor-info 时前任所占位符原样保留）
- P1 幂等（已有快照不重复预填）
- P3 原模板不变
- R1.3 联系方式区按"标签："独占段追加填充

对应 spec: b2-predecessor-communication-hardening Task 1/2
"""
from __future__ import annotations

import json
import shutil
import uuid
from datetime import date
from pathlib import Path
from unittest.mock import AsyncMock, MagicMock

import pytest
from docx import Document
from hypothesis import HealthCheck, given, settings
from hypothesis import strategies as st


# ---------------------------------------------------------------------------
# Helpers
# ---------------------------------------------------------------------------

def _make_b2_docx(tmp_path: Path, name: str = "b2_letter.docx") -> Path:
    doc = Document()
    doc.add_paragraph("【被审计单位名称】：")
    doc.add_paragraph("贵公司拟聘本所执行【20××】年度财务报表审计业务。")
    doc.add_paragraph("与前任事务所【前任会计师事务所的名称】进行沟通。")
    doc.add_paragraph("联系人：")
    doc.add_paragraph("地址：")
    fp = tmp_path / name
    doc.save(str(fp))
    return fp


def _extract_text(fp: Path) -> str:
    doc = Document(str(fp))
    return "\n".join(p.text for p in doc.paragraphs)


def _make_wp(file_path: str, wp_index_id=uuid.uuid4()) -> MagicMock:
    wp = MagicMock()
    wp.file_path = file_path
    wp.wp_index_id = wp_index_id
    return wp


def _make_db(
    *,
    client_name: str | None = "北京测试科技有限公司",
    audit_period_end: date | None = date(2025, 12, 31),
    wp_code: str = "B2-1",
    predecessor_info: dict | None = None,
) -> AsyncMock:
    db = AsyncMock()

    async def _execute(stmt, params=None):
        sql = str(getattr(stmt, "text", stmt))
        low = sql.lower()
        if "wp_index" in low and "wp_code" in low:
            r = MagicMock()
            r.scalar_one_or_none.return_value = wp_code
            return r
        if "select name, client_name, audit_period_end from projects" in low:
            r = MagicMock()
            if client_name or audit_period_end:
                row = MagicMock()
                row.__getitem__ = lambda self, i: [None, client_name, audit_period_end][i]
                r.first.return_value = row
            else:
                r.first.return_value = None
            return r
        if "b2-predecessor-info" in low:
            r = MagicMock()
            if predecessor_info is not None:
                r.first.return_value = (json.dumps(predecessor_info),)
            else:
                r.first.return_value = None
            return r
        # partner / staff / misstatement → None
        r = MagicMock()
        r.first.return_value = None
        r.scalar_one_or_none.return_value = None
        return r

    db.execute = _execute
    return db


def _cleanup(project_id: uuid.UUID, wp_code: str) -> None:
    snap = Path(f"storage/{project_id}/workpapers/{wp_code}.docx")
    if snap.parent.exists():
        shutil.rmtree(f"storage/{project_id}", ignore_errors=True)


# ---------------------------------------------------------------------------
# P4: 中文占位符替换正确
# ---------------------------------------------------------------------------

@pytest.mark.asyncio
@settings(max_examples=5, deadline=None,
          suppress_health_check=[HealthCheck.function_scoped_fixture])
@given(
    client=st.text(alphabet="测试科技有限公司北京上海ABC", min_size=2, max_size=10),
    year=st.integers(min_value=2015, max_value=2030),
    firm=st.text(alphabet="致同天健信永中和会计师事务所", min_size=2, max_size=12),
)
async def test_p4_cn_placeholders_replaced(tmp_path: Path, client, year, firm) -> None:
    """**Validates: R1.1, P4** — 三个中文占位符在有值时被正确替换。"""
    from app.routers.wp_editor_router import _prefill_word_template

    project_id = uuid.uuid4()
    wp_code = "B2-1"
    fp = _make_b2_docx(tmp_path, f"{uuid.uuid4().hex}.docx")
    wp = _make_wp(str(fp))
    db = _make_db(client_name=client, audit_period_end=date(year, 12, 31),
                  wp_code=wp_code, predecessor_info={"firmName": firm})
    try:
        result = await _prefill_word_template(db, wp, project_id, user=None)
        assert result is True
        text = _extract_text(Path(wp.file_path))
        assert client in text
        assert str(year) in text
        assert firm in text
        # 占位符已消失
        assert "【被审计单位名称】" not in text
        assert "【20××】" not in text
        assert "【前任会计师事务所的名称】" not in text
    finally:
        _cleanup(project_id, wp_code)


# ---------------------------------------------------------------------------
# P2: 空值保留
# ---------------------------------------------------------------------------

@pytest.mark.asyncio
async def test_p2_preserves_firm_placeholder_when_no_predecessor_info(tmp_path: Path) -> None:
    """**Validates: R1.2, P2** — 未录入前任信息时前任所占位符原样保留，实体名/年度仍替换。"""
    from app.routers.wp_editor_router import _prefill_word_template

    project_id = uuid.uuid4()
    wp_code = "B2-1"
    fp = _make_b2_docx(tmp_path, f"{uuid.uuid4().hex}.docx")
    wp = _make_wp(str(fp))
    db = _make_db(wp_code=wp_code, predecessor_info=None)
    try:
        result = await _prefill_word_template(db, wp, project_id, user=None)
        assert result is True
        text = _extract_text(Path(wp.file_path))
        assert "北京测试科技有限公司" in text
        assert "【前任会计师事务所的名称】" in text  # 保留
    finally:
        _cleanup(project_id, wp_code)


# ---------------------------------------------------------------------------
# R1.3: 联系方式区填充
# ---------------------------------------------------------------------------

@pytest.mark.asyncio
async def test_contact_labels_filled(tmp_path: Path) -> None:
    """**Validates: R1.3** — 独占整段的"联系人："/"地址："标签被追加值。"""
    from app.routers.wp_editor_router import _prefill_word_template

    project_id = uuid.uuid4()
    wp_code = "B2-1"
    fp = _make_b2_docx(tmp_path, f"{uuid.uuid4().hex}.docx")
    wp = _make_wp(str(fp))
    db = _make_db(wp_code=wp_code, predecessor_info={
        "firmName": "致同会计师事务所",
        "contactPerson": "王审计",
        "address": "北京市朝阳区",
    })
    try:
        result = await _prefill_word_template(db, wp, project_id, user=None)
        assert result is True
        text = _extract_text(Path(wp.file_path))
        assert "联系人：王审计" in text
        assert "地址：北京市朝阳区" in text
    finally:
        _cleanup(project_id, wp_code)


# ---------------------------------------------------------------------------
# P1: 幂等 + P3: 原模板不变
# ---------------------------------------------------------------------------

@pytest.mark.asyncio
async def test_p1_idempotent_and_p3_template_unchanged(tmp_path: Path) -> None:
    """**Validates: P1 幂等, P3 原模板不变** — 二次预填返回 False，原模板不含替换值。"""
    from app.routers.wp_editor_router import _prefill_word_template

    project_id = uuid.uuid4()
    wp_code = "B2-1"
    fp = _make_b2_docx(tmp_path, f"{uuid.uuid4().hex}.docx")
    original_text = _extract_text(fp)
    wp = _make_wp(str(fp))
    db = _make_db(wp_code=wp_code, predecessor_info={"firmName": "致同"})
    try:
        first = await _prefill_word_template(db, wp, project_id, user=None)
        assert first is True
        # 二次：快照已存在 → False
        wp2 = _make_wp(str(fp))
        second = await _prefill_word_template(db, wp2, project_id, user=None)
        assert second is False
        # 原模板文件未被修改（占位符仍在）
        assert _extract_text(fp) == original_text
        assert "【被审计单位名称】" in _extract_text(fp)
    finally:
        _cleanup(project_id, wp_code)


# ---------------------------------------------------------------------------
# 非 B2 底稿不查前任信息（回归：普通 docx 无 CN 占位符不受影响）
# ---------------------------------------------------------------------------

@pytest.mark.asyncio
async def test_non_b2_no_cn_lookup(tmp_path: Path) -> None:
    """**Validates: R1（边界）** — 无中文占位符的 docx 返回 False。"""
    from app.routers.wp_editor_router import _prefill_word_template

    project_id = uuid.uuid4()
    doc = Document()
    doc.add_paragraph("普通段落，无任何占位符。")
    fp = tmp_path / f"{uuid.uuid4().hex}.docx"
    doc.save(str(fp))
    wp = _make_wp(str(fp))
    db = _make_db(wp_code="A1", predecessor_info=None)
    try:
        result = await _prefill_word_template(db, wp, project_id, user=None)
        assert result is False
    finally:
        _cleanup(project_id, "A1")

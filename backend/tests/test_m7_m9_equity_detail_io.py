"""M7-2 / M9-2 明细表导入导出、M7 计提测试、M9 审定表行读取（spec m7-m9-service-modules-and-lazy-import-guard）。

真 SQLite + 真 SQL：测试里按真实 schema 建 ``working_paper``（service 用到的三列）与
``checklist_responses``（全部列 + ``(wp_id, item_id)`` 唯一），service 的语句原样执行；
端点用只挂两个 router 的最小应用 + httpx 真发请求。PostgreSQL 上的真实路由 + gate 见
``test_m7_m9_equity_detail_io_pg.py``。前后端契约（行字段 / 存储列 / 阈值）直接读前端源码比对。
"""

from __future__ import annotations

import io
import json
import math
import re
import uuid
from pathlib import Path
from types import SimpleNamespace

import pytest
import pytest_asyncio
import sqlalchemy as sa
from fastapi import HTTPException
from openpyxl import load_workbook
from sqlalchemy.ext.asyncio import AsyncSession, create_async_engine

from app.services import m7_special_reserve_service as m7
from app.services import m9_other_comprehensive_income_service as m9
from app.services import m_detail_sheet_io as core

_WP_DIR = Path(__file__).resolve().parents[2] / "audit-platform" / "frontend" / "src" / "components" / "workpaper"
_XLSX = "application/vnd.openxmlformats-officedocument.spreadsheetml.sheet"

_DDL = (
    "CREATE TABLE working_paper (id TEXT PRIMARY KEY, project_id TEXT NOT NULL, "
    "is_deleted BOOLEAN NOT NULL DEFAULT 0)",
    "CREATE TABLE checklist_responses (id TEXT PRIMARY KEY, project_id TEXT NOT NULL, wp_id TEXT NOT NULL, "
    "item_id TEXT NOT NULL, conclusion TEXT, remark TEXT, wp_ref TEXT, updated_by TEXT, "
    "created_at TIMESTAMP NOT NULL DEFAULT CURRENT_TIMESTAMP, "
    "updated_at TIMESTAMP NOT NULL DEFAULT CURRENT_TIMESTAMP, "
    "content_version INTEGER NOT NULL DEFAULT 1, UNIQUE (wp_id, item_id))",
)

#: 列序同 M7_DETAIL_SPEC：项目名称 / 期初 / 计提 / 费用化 / 资本化 / 期初AJE·RJE / 计提AJE·RJE /
#: 使用AJE·RJE / 上期期初·计提·使用·期末 / 备注
M7_ROW_A = ("安全生产费", 1000, 500, 200, 100, 10, 0, 5, 0, 3, 0, 800, 400, 200, 1000, "按收入计提")
M7_ROW_B = ("维简费", "1,234.50", "(100)", "-", None, None, None, None, None, None, None, None, None, None, None, None)
#: 列序同 M9_DETAIL_SPEC：OCI项目 / 分类 / 期初 / 税前 / 税额 / 贷方 / 借方 / 期初AJE·RJE / 税前AJE·RJE /
#: 税额AJE·RJE / 上期期初·税后净额·期末 / 来源底稿 / 备注
M9_ROW_A = ("其他权益工具投资公允价值变动", "不可重分类", 1000, 400, 100, 400, 100, 0, 0, 0, 0, 0, 0, 900, 50, 1000, "G8", None)
M9_ROW_B = ("外币财务报表折算差额", "以后能重分类进损益的其他综合收益", 0, 50, 0, 50, 0, 0, 0, 10, 0, 0, 0, 0, 0, 0, None, "备注")


@pytest_asyncio.fixture
async def db():
    engine = create_async_engine("sqlite+aiosqlite:///:memory:")
    async with engine.connect() as conn:
        for ddl in _DDL:
            await conn.execute(sa.text(ddl))
        await conn.commit()
        session = AsyncSession(bind=conn, expire_on_commit=False)
        try:
            yield session
        finally:
            await session.close()
    await engine.dispose()


@pytest_asyncio.fixture
async def wp(db):
    """一张底稿（id / project_id 都是带连字符的 UUID 字符串，与生产绑定形态一致）。"""
    ref = SimpleNamespace(id=str(uuid.uuid4()), project_id=str(uuid.uuid4()), user_id=str(uuid.uuid4()))
    await db.execute(
        sa.text("INSERT INTO working_paper (id, project_id, is_deleted) VALUES (:id, :pid, false)"),
        {"id": ref.id, "pid": ref.project_id},
    )
    await db.commit()
    return ref


async def _stored(db, wp_id: str) -> dict:
    result = await db.execute(
        sa.text(
            "SELECT item_id, conclusion, remark, updated_by, project_id "
            "FROM checklist_responses WHERE wp_id = :w"
        ),
        {"w": wp_id},
    )
    return {row.item_id: row for row in result.fetchall()}


async def _put(db, wp, item_id: str, *, conclusion=None, remark=None) -> None:
    await db.execute(
        sa.text(
            "INSERT INTO checklist_responses (id, project_id, wp_id, item_id, conclusion, remark) "
            "VALUES (:id, :pid, :wp, :item, :c, :r)"
        ),
        {"id": str(uuid.uuid4()), "pid": wp.project_id, "wp": wp.id, "item": item_id, "c": conclusion, "r": remark},
    )
    await db.commit()


def _fill(spec, rows, *, title_row: str | None = None) -> bytes:
    """在「导出模板」生成的文件上追加数据行（用户真实路径：先导出模板再填）。"""
    wb = load_workbook(io.BytesIO(core.export_template([spec]).getvalue()))
    ws = wb[spec.sheet_title]
    if title_row:
        ws.insert_rows(1)
        ws.cell(row=1, column=1, value=title_row)
    for row in rows:
        ws.append(list(row))
    buffer = io.BytesIO()
    wb.save(buffer)
    return buffer.getvalue()


def _sheet(content: bytes, title: str) -> list[list]:
    ws = load_workbook(io.BytesIO(content), data_only=True)[title]
    return [list(row) for row in ws.iter_rows(values_only=True)]


def _input_fields(spec) -> list[str]:
    return [col.field for col in spec.input_columns]


# ═══ 模板 ═══════════════════════════════════════════════════════════════════


@pytest.mark.parametrize("spec", [m7.M7_DETAIL_SPEC, m9.M9_DETAIL_SPEC], ids=["M7-2", "M9-2"])
def test_template_has_input_headers_and_guide(spec):
    wb = load_workbook(io.BytesIO(core.export_template([spec]).getvalue()))
    assert wb.sheetnames == [spec.sheet_title, core.GUIDE_SHEET_TITLE]
    header = [cell.value for cell in wb[spec.sheet_title][1]]
    assert header == [col.header for col in spec.input_columns]
    guide = "\n".join(str(r[0]) for r in wb[core.GUIDE_SHEET_TITLE].iter_rows(values_only=True) if r[0])
    assert "整表替换" in guide and spec.key_column.header in guide


# ═══ M7-2 导入：键、列、数值与界面一致 ═══════════════════════════════════════


async def test_m7_import_writes_the_keys_the_page_reads(db, wp):
    result = await m7.import_data(wp.id, _fill(m7.M7_DETAIL_SPEC, [M7_ROW_A, M7_ROW_B]), db, updated_by=wp.user_id)
    await db.commit()
    assert result == {"ok": True, "sheet": "M7-2", "imported_count": 2, "warnings": []}

    stored = await _stored(db, wp.id)
    assert set(stored) == {
        "M7-2-full-data", "M7-2-auditedEndTotal",
        "M7-2-row-0-data", "M7-2-row-0-auditedEnd", "M7-2-row-1-data", "M7-2-row-1-auditedEnd",
    }
    full = stored["M7-2-full-data"]
    assert full.conclusion is None, "界面从 remark 读 M7-2-full-data"
    assert {row.project_id for row in stored.values()} == {wp.project_id}
    assert {row.updated_by for row in stored.values()} == {wp.user_id}

    rows = json.loads(full.remark)
    assert [list(r) for r in rows] == [["key", *_input_fields(m7.M7_DETAIL_SPEC)]] * 2, "字段与顺序同前端行 JSON"
    assert '"itemName":"安全生产费"' in full.remark and '"beginning":1000,' in full.remark, "紧凑写法、整数不带 .0"
    assert rows[1] | {"key": None} == {
        "key": None, "itemName": "维简费", "beginning": 1234.5, "creditAccrual": -100,
        **{f: 0 for f in _input_fields(m7.M7_DETAIL_SPEC)[3:-1]}, "remark": "",
    }
    assert json.loads(stored["M7-2-row-0-data"].remark) == rows[0]
    # 审定期末：A = (1000+10+0) + (500+5+0) − (300+3+0) = 1212；B = 1234.5 − 100 = 1134.5
    assert stored["M7-2-row-0-auditedEnd"].remark == "1212"
    assert stored["M7-2-row-1-auditedEnd"].remark == "1134.5"
    assert stored["M7-2-auditedEndTotal"].remark == "2346.5"


async def test_m7_export_data_round_trips_through_import(db, wp):
    await m7.import_data(wp.id, _fill(m7.M7_DETAIL_SPEC, [M7_ROW_A, M7_ROW_B]), db)
    await db.commit()
    first = json.loads((await _stored(db, wp.id))["M7-2-full-data"].remark)

    exported = (await m7.generate_export_data(wp.id, db, sheet="M7-2")).getvalue()
    sheet = _sheet(exported, "明细表M7-2")
    header = sheet[0]
    assert header[: len(m7.M7_DETAIL_SPEC.input_columns)] == [c.header for c in m7.M7_DETAIL_SPEC.input_columns]
    assert "审定期末" in header and "变动率" in header
    by_header = [dict(zip(header, row)) for row in sheet[1:]]
    assert [r["项目名称"] for r in by_header] == ["安全生产费", "维简费", "合计"]
    assert by_header[0]["审定期末"] == 1212 and by_header[1]["审定期末"] == 1134.5
    assert by_header[2]["审定期末"] == 2346.5 and by_header[2]["期初余额"] == 2234.5
    assert math.isclose(by_header[0]["变动率"], 0.212)

    # 导出数据原样再导入：公式列忽略、合计行跳过，内容与第一次一致（只有生成的 key 不同）
    again = await m7.import_data(wp.id, exported, db)
    await db.commit()
    assert again["imported_count"] == 2
    assert again["warnings"] == ["已跳过合计 / 小计行 1 行（合计由系统按明细重算）"]
    second = json.loads((await _stored(db, wp.id))["M7-2-full-data"].remark)
    assert [r | {"key": None} for r in second] == [r | {"key": None} for r in first]


async def test_reimport_with_fewer_rows_clears_stale_row_keys(db, wp):
    rows = [M7_ROW_A, ("其他", 1, 0, 0, 0, 0, 0, 0, 0, 0, 0, 0, 0, 0, 0, None), M7_ROW_B]
    await m7.import_data(wp.id, _fill(m7.M7_DETAIL_SPEC, rows), db)
    await _put(db, wp, "M7-2-row-9-note", remark="同前缀但不是逐行派生键，不得动")
    await m7.import_data(wp.id, _fill(m7.M7_DETAIL_SPEC, [M7_ROW_B]), db)
    await db.commit()

    stored = await _stored(db, wp.id)
    assert [r["itemName"] for r in json.loads(stored["M7-2-full-data"].remark)] == ["维简费"]
    assert stored["M7-2-row-0-auditedEnd"].remark == "1134.5"
    for stale in ("M7-2-row-1-data", "M7-2-row-1-auditedEnd", "M7-2-row-2-data", "M7-2-row-2-auditedEnd"):
        assert stored[stale].remark is None, f"{stale} 应被清空"
    assert stored["M7-2-row-9-note"].remark == "同前缀但不是逐行派生键，不得动"
    assert stored["M7-2-auditedEndTotal"].remark == "1134.5"


async def test_import_only_touches_the_storage_column(db, wp):
    await _put(db, wp, "M7-2-full-data", conclusion="界面别处写的结论", remark="[]")
    await m7.import_data(wp.id, _fill(m7.M7_DETAIL_SPEC, [M7_ROW_A]), db)
    await db.commit()
    stored = (await _stored(db, wp.id))["M7-2-full-data"]
    assert stored.conclusion == "界面别处写的结论"
    assert json.loads(stored.remark)[0]["itemName"] == "安全生产费"


async def test_header_may_sit_below_a_title_row(db, wp):
    content = _fill(m7.M7_DETAIL_SPEC, [M7_ROW_A], title_row="某公司 专项储备明细表")
    assert (await m7.import_data(wp.id, content, db))["imported_count"] == 1


def _without_column(spec, header: str, rows) -> bytes:
    content = _fill(spec, rows)
    wb = load_workbook(io.BytesIO(content))
    ws = wb[spec.sheet_title]
    idx = [c.value for c in ws[1]].index(header) + 1
    ws.delete_cols(idx)
    buffer = io.BytesIO()
    wb.save(buffer)
    return buffer.getvalue()


def _with_duplicate_header(spec) -> bytes:
    content = _fill(spec, [M7_ROW_A])
    wb = load_workbook(io.BytesIO(content))
    ws = wb[spec.sheet_title]
    ws.cell(row=1, column=ws.max_column + 1, value="期初 余额")  # 归一后与「期初余额」相同
    buffer = io.BytesIO()
    wb.save(buffer)
    return buffer.getvalue()


_BAD_FILES = {
    "empty": (lambda: b"", "上传的文件为空"),
    "not_xlsx": (lambda: b"not a workbook", "无法解析文件"),
    "missing_column": (lambda: _without_column(m7.M7_DETAIL_SPEC, "备注", [M7_ROW_A]), "缺少列：备注"),
    "duplicate_header": (lambda: _with_duplicate_header(m7.M7_DETAIL_SPEC), "表头重复：期初余额"),
    "bad_number": (lambda: _fill(m7.M7_DETAIL_SPEC, [("甲", "abc")]), "第2行「期初余额」不是有效数字：abc"),
    "boolean_amount": (lambda: _fill(m7.M7_DETAIL_SPEC, [("甲", True)]), "第2行「期初余额」不是有效数字"),
    "empty_name": (lambda: _fill(m7.M7_DETAIL_SPEC, [(None, 100)]), "第2行「项目名称」为空"),
    "only_total": (lambda: _fill(m7.M7_DETAIL_SPEC, [("合计", 100)]), "没有可导入的数据行"),
    "too_many_rows": (lambda: _fill(m7.M7_DETAIL_SPEC, [(f"项{i}", i) for i in range(501)]), "超过 500 行上限"),
}


@pytest.mark.parametrize("case", sorted(_BAD_FILES))
async def test_bad_file_is_rejected_and_nothing_is_written(db, wp, case):
    make, message = _BAD_FILES[case]
    with pytest.raises(ValueError, match=re.escape(message)):
        await m7.import_data(wp.id, make(), db)
    assert await _stored(db, wp.id) == {}


@pytest.mark.parametrize("service, unknown", [(m7, "M7-4"), (m9, "M9-4")], ids=["M7", "M9"])
async def test_unknown_sheet_is_a_readable_400_on_all_three_states(db, wp, service, unknown):
    calls = (
        service.generate_export_template(wp.id, db, sheet=unknown),
        service.generate_export_data(wp.id, db, sheet=unknown),
        service.import_data(wp.id, b"irrelevant", db, sheet=unknown),
    )
    for call in calls:
        with pytest.raises(HTTPException) as caught:
            await call
        assert caught.value.status_code == 400
        assert caught.value.detail == f"不支持的sheet: {unknown}。支持: {sorted(service.SUPPORTED_SHEETS)}"


async def test_import_requires_an_existing_live_workpaper(db, wp):
    content = _fill(m7.M7_DETAIL_SPEC, [M7_ROW_A])
    await db.execute(sa.text("INSERT INTO working_paper (id, project_id, is_deleted) VALUES (:id, :p, true)"),
                     {"id": (deleted := str(uuid.uuid4())), "p": wp.project_id})
    for wp_id in ("not-a-uuid", str(uuid.uuid4()), deleted):
        with pytest.raises(HTTPException) as caught:
            await m7.import_data(wp_id, content, db)
        assert (caught.value.status_code, caught.value.detail) == (404, core.NOT_FOUND_DETAIL)
    with pytest.raises(HTTPException):
        await m7.generate_export_data("not-a-uuid", db)


async def test_export_data_reads_back_in_the_page_order(db, wp):
    # 没有数据：只有表头，「编制说明」写明
    empty = (await m7.generate_export_data(wp.id, db)).getvalue()
    assert len(_sheet(empty, "明细表M7-2")) == 1
    assert any("暂无已保存的明细行" in str(r[0]) for r in _sheet(empty, core.GUIDE_SHEET_TITLE))

    # 整表键坏 JSON ⇒ 落到旧键 M7-2-rows（同界面 _restoreRowsFromResponses），并在说明里写明
    await _put(db, wp, "M7-2-full-data", remark="{坏的")
    await _put(db, wp, "M7-2-rows", remark=json.dumps([{"itemName": "旧版行", "beginning": "88"}]))
    exported = (await m7.generate_export_data(wp.id, db)).getvalue()
    assert [row[0] for row in _sheet(exported, "明细表M7-2")[1:]] == ["旧版行", "合计"]
    assert _sheet(exported, "明细表M7-2")[1][1] == 88
    assert any("M7-2-full-data" in str(r[0]) and "JSON" in str(r[0]) for r in _sheet(exported, core.GUIDE_SHEET_TITLE))


async def test_empty_full_data_falls_back_like_the_page(db, wp):
    await _put(db, wp, "M7-2-full-data", remark="[]")
    await _put(db, wp, "M7-2-rows", remark=json.dumps([{"itemName": "旧版行"}]))
    rows, warnings = await core.load_rows(db, wp.id, m7.M7_DETAIL_SPEC)
    assert [r["itemName"] for r in rows] == ["旧版行"] and warnings == []


# ═══ M9-2 ═══════════════════════════════════════════════════════════════════


async def test_m9_import_writes_conclusion_and_maps_categories(db, wp):
    rows = [M9_ROW_A, M9_ROW_B, ("现金流量套期损益", "reclass", 0, 0, 0, 0, 0)]
    result = await m9.import_data(wp.id, _fill(m9.M9_DETAIL_SPEC, rows), db, updated_by=wp.user_id)
    await db.commit()
    assert result["imported_count"] == 3

    stored = await _stored(db, wp.id)
    full = stored["M9-2-detail-rows"]
    assert full.remark is None, "界面经 getField 读 M9-2-detail-rows 的 conclusion"
    parsed = json.loads(full.conclusion)
    assert [list(r) for r in parsed] == [["key", *_input_fields(m9.M9_DETAIL_SPEC)]] * 3
    assert [r["ociCategory"] for r in parsed] == ["nonReclass", "reclass", "reclass"]
    assert parsed[0]["sourceWpCode"] == "G8" and parsed[1]["remark"] == "备注"
    # 审定期末：A = 1000 + (400 − 100) = 1300；B = 0 + ((50 + 10) − 0) = 60
    assert [stored[f"M9-2-row-{i}-auditedEnd"].remark for i in range(3)] == ["1300", "60", "0"]
    assert stored["M9-2-auditedEndTotal"].remark == "1360"
    assert stored["M9-2-row-1-data"].conclusion is None and json.loads(stored["M9-2-row-1-data"].remark) == parsed[1]

    exported = _sheet((await m9.generate_export_data(wp.id, db)).getvalue(), "明细表M9-2")
    by_header = [dict(zip(exported[0], row)) for row in exported[1:]]
    assert [r["分类"] for r in by_header] == ["不可重分类", "可重分类", "可重分类", None]
    assert [r["审定期末"] for r in by_header] == [1300, 60, 0, 1360]


@pytest.mark.parametrize("category", ["其他", None])
async def test_m9_rejects_unknown_or_blank_category(db, wp, category):
    content = _fill(m9.M9_DETAIL_SPEC, [("某项目", category, 1)])
    with pytest.raises(ValueError, match="第2行「分类」应填「不可重分类」或「可重分类」"):
        await m9.import_data(wp.id, content, db)
    assert await _stored(db, wp.id) == {}


# ═══ 数字 ═══════════════════════════════════════════════════════════════════


@pytest.mark.parametrize(
    "raw, expected",
    [(None, 0.0), ("", 0.0), ("  ", 0.0), ("-", 0.0), ("—", 0.0), (12, 12.0), (3.5, 3.5),
     ("1,234.50", 1234.5), ("¥ 2,000", 2000.0), ("(300)", -300.0), ("（12.5）", -12.5), ("-7", -7.0), ("1e3", 1000.0)],
)
def test_parse_amount_accepts_accounting_notation(raw, expected):
    assert core.parse_amount(raw) == expected


@pytest.mark.parametrize("raw", ["abc", "12元", "5%", True, float("nan"), float("inf"), "1,2,3.4.5"])
def test_parse_amount_rejects_non_numbers(raw):
    with pytest.raises(ValueError):
        core.parse_amount(raw)


@pytest.mark.parametrize(
    "value, text",
    [(1500.0, "1500"), (1500.5, "1500.5"), (-0.0, "0"), (0.1 + 0.2, "0.30000000000000004"),
     (1e-6, "0.000001"), (1e-7, "1e-7"), (1.5e-7, "1.5e-7"), (1e16, "10000000000000000"), (1e21, "1e+21")],
)
def test_js_number_matches_javascript_string(value, text):
    """与前端 ``String(number)`` 逐字一致（表中右侧为 Node / V8 实测输出）。"""
    assert core.js_number(value) == text


# ═══ M7-4 计提测试 ═════════════════════════════════════════════════════════


def _accrual(method="output", tiers=(), revenue=0.0, revenue_rate=0.0, booked=0.0):
    return m7.run_accrual_test(
        method=method, tiers=[{"output": o, "rate": r} for o, r in tiers],
        revenue=revenue, revenue_rate=revenue_rate, booked_accrual=booked,
    )


def test_accrual_by_output_sums_tiers():
    result = _accrual(tiers=[(100, 5), (50, 4)], booked=600)
    assert set(result) == {"method", "estimated", "booked", "diff", "is_material", "detail"}
    assert (result["estimated"], result["booked"], result["diff"], result["is_material"]) == (700, 600, 100, False)
    assert [t["amount"] for t in result["detail"]["tiers"]] == [500, 200]
    assert result["detail"]["conclusion"] == "少计提 100.00 元，未超过重要性阈值。"


def test_accrual_by_revenue_and_materiality():
    result = _accrual("REVENUE ", revenue=10_000_000, revenue_rate=0.015, booked=120_000)
    assert (result["method"], result["estimated"], result["diff"], result["is_material"]) == ("revenue", 150_000, 30_000, True)
    assert "超过重要性阈值" in result["detail"]["conclusion"]


@pytest.mark.parametrize(
    "diff, estimated, material",
    [(1000, 1, False), (1000.01, 1_000_000, False), (-1500, 0, True), (5000, 100_000, False), (5001, 100_000, True)],
)
def test_materiality_follows_the_page_rule(diff, estimated, material):
    """同 ``useM7AccrualTest.isDiffExceedThreshold``：|差异| ≤ 1000 否；应计提为 0 看绝对额；否则相对差异 > 5%。"""
    assert m7.is_accrual_diff_material(diff, estimated) is material


@pytest.mark.parametrize(
    "kwargs, message",
    [({"method": "按产量"}, "不支持的计提方式"),
     ({"method": "revenue", "revenue": float("inf")}, "「营业收入」必须是有限数字"),
     ({"booked": float("nan")}, "「账面已计提金额」必须是有限数字"),
     ({"tiers": [(float("inf"), 1)]}, "「第1档产量」必须是有限数字")],
)
def test_accrual_rejects_bad_input(kwargs, message):
    with pytest.raises(ValueError, match=message):
        _accrual(**kwargs)


# ═══ M9-1 审定表行（公式校验取数）════════════════════════════════════════════


async def test_adjudication_rows_follow_the_page_restore_rule(db, wp):
    row = {"key": "m9-adj-1", "itemName": "外币报表折算差额", "beginning": 100, "creditAmount": "30",
           "debitAmount": 10, "unadjusted": 125}
    await _put(db, wp, "M9-1-m9-adj-1-data", remark=json.dumps(row))
    await _put(db, wp, "M9-1-row-0-data", remark=json.dumps({"itemName": "旧格式行", "beginning": 5}))
    await _put(db, wp, "M9-1-m9-adj-2-data", remark="{坏的")
    await _put(db, wp, "M9-1-m9-adj-3-data", remark=json.dumps([1, 2]))
    await _put(db, wp, "M9-1-auditConclusion", remark="结论文字，不是行")
    await _put(db, wp, "M9-1-m9-adj-1-name", remark="派生键，不是行")
    await _put(db, wp, "M9-1-m9-adj-4-data", conclusion="只有 conclusion 的不算（界面只读 remark）")

    rows, warnings = await m9.get_adjudication_rows(wp.id, db)
    assert rows == [
        {"rowKey": "m9-adj-1", "itemName": "外币报表折算差额", "begin": 100.0, "credit": 30.0, "debit": 10.0,
         "reportedEnd": 125.0},
        {"rowKey": "M9-1-row-0-data", "itemName": "旧格式行", "begin": 5.0, "credit": 0.0, "debit": 0.0,
         "reportedEnd": None},
    ]
    assert len(warnings) == 2 and "M9-1-m9-adj-2-data" in warnings[0] and "M9-1-m9-adj-3-data" in warnings[1]


# ═══ 端点（只挂两个 router 的最小应用；gate 见文末与 PG 用例）═══════════════════


@pytest_asyncio.fixture
async def client(db, wp):
    from fastapi import FastAPI
    from httpx import ASGITransport, AsyncClient

    from app.core.database import get_db
    from app.deps import get_current_user
    from app.routers.m7_special_reserve import router as m7_router
    from app.routers.m9_other_comprehensive_income import router as m9_router

    app = FastAPI()
    app.include_router(m7_router)
    app.include_router(m9_router)

    async def _db():
        yield db

    app.dependency_overrides[get_db] = _db
    app.dependency_overrides[get_current_user] = lambda: SimpleNamespace(id=uuid.UUID(wp.user_id))
    async with AsyncClient(transport=ASGITransport(app=app), base_url="http://test") as http:
        yield http


async def test_m7_endpoints_end_to_end(client, db, wp):
    base = f"/api/m7-special-reserve/{wp.id}"
    template = await client.get(f"{base}/export-template", params={"sheet": "M7-2"})
    assert template.status_code == 200 and template.headers["content-type"] == _XLSX
    assert "filename*=UTF-8''" in template.headers["content-disposition"]
    assert load_workbook(io.BytesIO(template.content)).sheetnames == ["明细表M7-2", core.GUIDE_SHEET_TITLE]

    bad_sheet = await client.get(f"{base}/export-data", params={"sheet": "M7-9"})
    assert bad_sheet.status_code == 400 and "不支持的sheet: M7-9" in bad_sheet.json()["detail"]

    # 与前端 useM7ImportExport 同形：sheet 放在表单字段里（路由按 query 收不到 ⇒ 缺省 M7-2）
    content = _fill(m7.M7_DETAIL_SPEC, [M7_ROW_A])
    imported = await client.post(
        f"{base}/import-data", files={"file": ("m7.xlsx", content, _XLSX)}, data={"sheet": "M7-2", "wp_id": wp.id}
    )
    assert imported.status_code == 200, imported.text
    assert imported.json()["imported_count"] == 1
    assert (await _stored(db, wp.id))["M7-2-full-data"].updated_by == wp.user_id

    data = await client.get(f"{base}/export-data", params={"sheet": "M7-2"})
    assert [row[0] for row in _sheet(data.content, "明细表M7-2")[1:]] == ["安全生产费", "合计"]

    garbage = await client.post(f"{base}/import-data", files={"file": ("m7.xlsx", b"zzz", _XLSX)})
    assert garbage.status_code == 400 and "无法解析文件" in garbage.json()["detail"]

    accrual = await client.post(
        f"{base}/accrual-test", json={"method": "revenue", "revenue": 100000, "revenue_rate": 0.02, "booked_accrual": 500}
    )
    assert accrual.status_code == 200
    assert accrual.json() | {"detail": None} == {
        "method": "revenue", "estimated": 2000, "booked": 500, "diff": 1500, "is_material": True, "detail": None,
    }
    wrong = await client.post(f"{base}/accrual-test", json={"method": "monthly"})
    assert wrong.status_code == 400 and "不支持的计提方式" in wrong.json()["detail"]


async def test_m9_endpoints_end_to_end(client, db, wp):
    base = f"/api/m9-other-comprehensive-income/{wp.id}"
    content = _fill(m9.M9_DETAIL_SPEC, [M9_ROW_A])
    imported = await client.post(f"{base}/import-data", files={"file": ("m9.xlsx", content, _XLSX)})
    assert imported.status_code == 200 and imported.json()["imported_count"] == 1
    data = await client.get(f"{base}/export-data")
    assert [row[0] for row in _sheet(data.content, "明细表M9-2")[1:]] == ["其他权益工具投资公允价值变动", "合计"]

    def _adj(key, name, unadjusted):
        return json.dumps({"key": key, "itemName": name, "beginning": 100, "creditAmount": 30,
                           "debitAmount": 10, "unadjusted": unadjusted, "aje": 0, "rje": 0})

    await _put(db, wp, "M9-1-a-data", remark=_adj("a", "相符", 120))
    await _put(db, wp, "M9-1-b-data", remark=_adj("b", "不符", 200))
    await _put(db, wp, "M9-1-c-data", remark="{坏的")
    checked = await client.get(f"{base}/validate-formulas")
    assert checked.status_code == 200, checked.text
    body = checked.json()
    assert (body["total_rows"], body["valid_rows"], body["invalid_rows"]) == (2, 1, 1)
    assert [
        (r["row_key"], r["item_name"], r["calculated_end"], r["reported_end"], r["diff"], r["is_valid"])
        for r in body["results"]
    ] == [("a", "相符", 120, 120, 0, True), ("b", "不符", 120, 200, -80, False)]
    assert body["warnings"] == ["「M9-1-c-data」的内容不是有效的 JSON，未参与校验"]


def test_real_app_mounts_m7_m9_routes_behind_the_wp_gate():
    """生产注册：两个 router 的全部 {wp_id} 路由都挂 router 级 dedicated_wp_gate（读 / 写按方法分类）。"""
    from app.main import app
    from app.routers._wp_gate import dedicated_wp_gate

    found = {}
    for route in app.routes:
        path = getattr(route, "path", "")
        if path.startswith(("/api/m7-special-reserve/{wp_id}/", "/api/m9-other-comprehensive-income/{wp_id}/")):
            found[path] = {dep.call for dep in route.dependant.dependencies}
    expected_suffixes = {"export-template", "export-data", "import-data", "accrual-test", "validate-formulas", "tb-writeback"}
    assert {p.rsplit("/", 1)[1] for p in found} == expected_suffixes
    assert all(dedicated_wp_gate in calls for calls in found.values()), sorted(
        p for p, calls in found.items() if dedicated_wp_gate not in calls
    )


# ═══ 与前端源码的契约（行字段 / 存储列 / 读回顺序 / 阈值；改一侧另一侧打红）═══════


def _ts(rel: str) -> str:
    """读前端源码并去掉注释 —— 注释里写的不算数。"""
    text = (_WP_DIR / rel).read_text(encoding="utf-8")
    text = re.sub(r"/\*.*?\*/", "", text, flags=re.S)
    return re.sub(r"(?m)^\s*//.*$", "", text)


def _row_fields(source: str, anchor: str, var: str = "row") -> list[str]:
    """``anchor`` 之后第一个对象字面量里 ``field: {var}.xxx`` 的字段名（按出现顺序）。"""
    start = source.index(anchor)
    body = source[start : source.index("})", start)]
    return re.findall(rf"(\w+):\s*{var}\.\w+", body)


@pytest.mark.parametrize(
    "spec, composable, row_anchor, full_anchor, full_column",
    [
        (m7.M7_DETAIL_SPEC, "composables/useM7Detail.ts", "debouncedSave(`M7-2-row-${rowIndex}-data`",
         "debouncedSave('M7-2-full-data'", "remark"),
        (m9.M9_DETAIL_SPEC, "composables/useM9Detail.ts", "debouncedSave(`M9-2-row-${rowIndex}-data`",
         "debouncedSave('M9-2-detail-rows'", "conclusion"),
    ],
    ids=["M7-2", "M9-2"],
)
def test_backend_writes_what_the_page_writes(spec, composable, row_anchor, full_anchor, full_column):
    src = _ts(composable)
    assert _row_fields(src, row_anchor) == ["key", *_input_fields(spec)], "逐行 JSON 字段与顺序"
    assert re.search(rf"{re.escape(row_anchor)},\s*\{{\s*remark:", src)
    assert re.search(rf"{re.escape(full_anchor)},\s*\{{\s*{full_column}:", src), f"整表键存 {full_column}"
    assert spec.full_storage == full_column
    for derived in (f"`{spec.row_item_prefix}${{rowIndex}}-auditedEnd`", f"'{spec.total_item_id}'"):
        assert re.search(rf"debouncedSave\({re.escape(derived)},\s*\{{\s*remark:", src), derived
    if spec is m9.M9_DETAIL_SPEC:  # M9 的整表键字段单独列在 _syncFullData 里
        assert _row_fields(src, full_anchor) == ["key", *_input_fields(spec)]


def test_pages_read_back_in_the_order_the_backend_reads():
    m7_page = _ts("m7/core/M7TabDetail.vue")
    restore = m7_page[m7_page.index("function _restoreRowsFromResponses") :]
    assert restore.index("get('M7-2-full-data')") < restore.index("get('M7-2-rows')")
    assert "fullDataResp?.remark" in restore and "rowsResp?.remark" in restore
    assert [src for src, _col in m7.M7_DETAIL_SPEC.fallback_sources] == ["M7-2-rows"]
    to_row = m7_page[m7_page.index("function _toDetailRow") :]
    assert set(_input_fields(m7.M7_DETAIL_SPEC)) <= set(re.findall(r"\br\.(\w+)", to_row[: to_row.index("\n}")]))

    m9_page = _ts("m9/core/M9TabDetail.vue")
    assert "formData.getField('2', 'detail-rows')" in m9_page
    init = m9_page[m9_page.index("function _initDefaultRows") :]
    assert set(_input_fields(m9.M9_DETAIL_SPEC)) <= set(re.findall(r"\bitem\.(\w+)", init))
    get_field = _ts("composables/useM9FormData.ts")
    get_field = get_field[get_field.index("function getField") :]
    assert "resp.conclusion" in get_field[: get_field.index("\n  }")]


def test_adjudication_rows_read_the_fields_the_page_writes():
    src = _ts("composables/useM9Adjudication.ts")
    fields = _row_fields(src, "debouncedSave(`M9-1-${k}-data`")
    assert {"key", "itemName", "beginning", "creditAmount", "debitAmount", "unadjusted"} <= set(fields)
    assert "key.startsWith('M9-1-') && key.endsWith('-data') && resp.remark" in _ts("m9/core/M9TabAdjudication.vue")
    assert (m9.ADJUDICATION_ITEM_PREFIX, m9.ADJUDICATION_ITEM_SUFFIX) == ("M9-1-", "-data")


def test_accrual_threshold_and_rule_match_the_page():
    src = _ts("composables/useM7AccrualTest.ts")
    block = re.search(r"const DEFAULT_THRESHOLD[^=]*=\s*\{(?P<body>[^}]*)\}", src)
    assert block, "前端 DEFAULT_THRESHOLD 不见了"
    absolute = float(re.search(r"absolute:\s*([\d.]+)", block["body"]).group(1))
    relative = float(re.search(r"relative:\s*([\d.]+)", block["body"]).group(1))
    assert (absolute, relative) == (m7.ACCRUAL_THRESHOLD_ABSOLUTE, m7.ACCRUAL_THRESHOLD_RELATIVE)
    rule = src[src.index("function isDiffExceedThreshold") :]
    rule = rule[: rule.index("\n  }")]
    for snippet in (
        "if (Math.abs(diffValue) <= threshold.absolute) return false",
        "if (estimatedValue === 0) return Math.abs(diffValue) > threshold.absolute",
        "return Math.abs(diffValue / estimatedValue) > threshold.relative",
    ):
        assert snippet in rule, snippet


def test_import_keys_do_not_collide_with_formula_push_owned_keys():
    """推送引擎独占的键会被 checklist 保存端点剔除；导入写的键不得落在里面。"""
    from app.services.formula_push.owned_keys import owned_item_ids

    for code, spec in (("M7", m7.M7_DETAIL_SPEC), ("M9", m9.M9_DETAIL_SPEC)):
        written = {spec.full_item_id, spec.total_item_id} | {
            f"{spec.row_item_prefix}{i}-{suffix}" for i in range(500) for suffix in ("data", "auditedEnd")
        }
        assert not (owned_item_ids(code) & written), code

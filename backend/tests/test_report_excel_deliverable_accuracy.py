"""财务报表交付导出准确性（spec deliverable-trio-generation-accuracy Phase A）。

真实模板 soe_standalone.xlsx + SQLite 夹具，钉住 2026-10-16 真实数据实测出的缺陷：
1. 库内通用编号行按行名落到模板行（不按编码错格）；
2. 模板「实收资本 = SUM(构成明细)」而库内只有总额 ⇒ 写总额（旧导出整行丢 1.5 亿）；
3. 单体口径「*2.少数股东损益」模板公式清空（旧导出把全部净利润算成少数股东损益）；
4. 附注列按交付件显示编号（单体国企「八」显示为「七」，无数据的节不占号）；
5. 减值表别名 asset_impairment、期间报表表头「YYYY年度」、导出提醒。
"""

from __future__ import annotations

import uuid
from datetime import date, datetime, timezone
from decimal import Decimal
from io import BytesIO

import pytest
import pytest_asyncio
from openpyxl import Workbook, load_workbook
from sqlalchemy import MetaData
from sqlalchemy.dialects.sqlite.base import SQLiteTypeCompiler
from sqlalchemy.ext.asyncio import AsyncSession, async_sessionmaker, create_async_engine

from app.models.core import Project
from app.models.report_models import DisclosureNote, FinancialReport, FinancialReportType

SQLiteTypeCompiler.visit_JSONB = SQLiteTypeCompiler.visit_JSON
if not hasattr(SQLiteTypeCompiler, "visit_ARRAY"):
    SQLiteTypeCompiler.visit_ARRAY = lambda self, type_, **kw: "TEXT"

from app.services.note_deliverable_outline import FsNoteRefPlan, NoteTemplateOutline, OutlineRegion  # noqa: E402
from app.services.report_excel_exporter import ReportExcelExporter  # noqa: E402
from app.services.report_excel_sheet_ops import (  # noqa: E402
    blank_consolidated_only_formulas,
    fill_formula_rows_without_inputs,
)
from app.services.template_manifest_loader import resolve_template_base_dir  # noqa: E402

_REAL = resolve_template_base_dir() / "financial_statements" / "soe_standalone.xlsx"
_TABLES = [Project.__table__, FinancialReport.__table__, DisclosureNote.__table__]
YEAR = 2024


@pytest_asyncio.fixture
async def db():
    engine = create_async_engine("sqlite+aiosqlite:///:memory:")
    async with engine.begin() as conn:
        await conn.run_sync(lambda c: MetaData().create_all(c, tables=_TABLES))
    factory = async_sessionmaker(engine, class_=AsyncSession, expire_on_commit=False)
    async with factory() as session:
        yield session
    await engine.dispose()


def _fr(pid, rt, code, name, cur=None, prior=None, formula=None):
    return FinancialReport(
        id=uuid.uuid4(), project_id=pid, year=YEAR, report_type=rt, row_code=code, row_name=name,
        current_period_amount=None if cur is None else Decimal(str(cur)),
        prior_period_amount=None if prior is None else Decimal(str(prior)),
        formula_used=formula, indent_level=1, is_total_row=False, generated_at=datetime.now(timezone.utc),
    )


def _note(pid, section, title, text):
    return DisclosureNote(
        id=uuid.uuid4(), project_id=pid, year=YEAR, note_section=section, section_title=title,
        text_content=text, status="draft", is_deleted=False, is_empty=False,
    )


@pytest_asyncio.fixture
async def project(db: AsyncSession):
    pid = uuid.uuid4()
    db.add(Project(id=pid, name="测试单体_2024", client_name="测试单体有限公司", template_type="soe",
                   report_scope="standalone", audit_period_end=date(YEAR, 12, 31)))
    bs, is_, imp = (FinancialReportType.balance_sheet, FinancialReportType.income_statement,
                    FinancialReportType.impairment_provision)
    db.add_all([
        # 资产负债表用库内「通用编号」：BS-004 衍生金融资产、BS-006 应收账款（模板分别是 BS-006 / BS-008）
        _fr(pid, bs, "BS-002", "货币资金", 100, 80, "TB('1001','期末余额')"),
        _fr(pid, bs, "BS-004", "衍生金融资产", 9, 0, "TB('1102','期末余额')"),
        _fr(pid, bs, "BS-006", "应收账款", 500, 400, "TB('1122','期末余额')"),
        _fr(pid, bs, "BS-081", "实收资本（或股本）", 150, 150, "TB('4001','期末余额')"),
        _fr(pid, bs, "BS-089", "所有者权益合计", 150, 150, "ROW('BS-081')"),
        _fr(pid, is_, "IS-001", "营业收入", 1000, None, "TB('6001','本期发生额')"),
        _fr(pid, is_, "IS-024", "四、净利润", 30, None, "ROW('IS-022') - ROW('IS-023')"),
        _fr(pid, imp, "IMP-001", "一、坏账准备", 20, 10, "TB('1231','期末余额')"),
    ])
    db.add_all([
        _note(pid, "八、1", "货币资金", "期末货币资金均为银行存款。"),
        _note(pid, "八、5", "应收账款", "应收账款按账龄披露如下。"),
    ])
    await db.flush()
    return pid


def _sheet(wb, *keywords):
    for name in wb.sheetnames:
        if all(k in name for k in keywords):
            return wb[name]
    raise AssertionError(f"缺 sheet {keywords}: {wb.sheetnames}")


@pytest.mark.skipif(not _REAL.is_file(), reason="报表模板缺失")
@pytest.mark.asyncio
async def test_real_template_export_accuracy(db, project):
    exporter = ReportExcelExporter(db)
    out = await exporter.export(
        project, YEAR, mode="audited",
        report_types=["balance_sheet", "income_statement", "asset_impairment"],
        flatten_formulas=True,
    )
    wb = load_workbook(BytesIO(out.getvalue()), data_only=True)

    bs = _sheet(wb, "资产负债表")
    assert "货币资金" in str(bs["A6"].value) and bs["C6"].value == pytest.approx(100)
    assert "应收账款" in str(bs["A12"].value) and bs["C12"].value == pytest.approx(500)
    assert "衍生金融资产" in str(bs["A10"].value) and bs["C10"].value == pytest.approx(9)
    # 附注列：只有 八、1 / 八、5 有附注 ⇒ 交付件里是「七、1」「七、2」；衍生金融资产无附注留空
    assert bs["B6"].value == "七、1"
    assert bs["B12"].value == "七、2"
    assert bs["B10"].value in (None, "")

    bs2 = _sheet(wb, "资产负债表", "续")
    assert "实收资本" in str(bs2["A53"].value)
    assert bs2["C53"].value == pytest.approx(150)  # 模板 =SUM(国家资本…外商资本)，库内只有总额
    assert bs2["D53"].value == pytest.approx(150)

    is_ws = _sheet(wb, "利润表")
    assert is_ws["A2"].value == f"{YEAR}年度"
    assert "少数股东损益" in str(is_ws["A53"].value)
    assert is_ws["C53"].value in (None, "")  # 单体：带 * 合并专用行不按公式算成净利润

    _sheet(wb, "资产减值准备")  # asset_impairment 别名 → 减值表保留

    warnings = "\n".join(exporter.export_warnings)
    assert "衍生金融资产" in warnings and "附注列" in warnings


def test_compound_note_ref_cell_renders_both_parts_and_drops_empty_separator():
    ol = NoteTemplateOutline(
        chapters=["财务报表主要项目注释", "母公司财务报表的主要项目附注"],
        regions=[
            OutlineRegion(0, "八、1", "货币资金", "货币资金", 1, "Heading 3"),
            OutlineRegion(1, "八、5", "应收账款", "应收账款", 1, "Heading 3"),
            OutlineRegion(2, "十二、应收账款", "应收账款", "应收账款", 2, "Heading 3"),
        ],
        chapter_pos_of_prefix={"八": 1, "十二": 2},
    )
    plan = FsNoteRefPlan(outline=ol, numbers={0: "八、1", 1: "八、2", 2: "十二、1"}, has_notes=True)
    wb = Workbook()
    ws = wb.active
    ws["A5"], ws["B5"] = "货币资金", "{{note_ref:BS-002}};{{note_ref:BS-002:parent}}"
    ws["A6"], ws["B6"] = "应收账款", "{{note_ref:BS-008}};{{note_ref:BS-008:parent}}"
    exporter = ReportExcelExporter(db=None)  # type: ignore[arg-type]
    exporter._fill_by_placeholders(ws, {}, {}, True, {}, None, {}, note_plan=plan)
    assert ws["B5"].value == "八、1"
    assert ws["B6"].value == "八、2;十二、1"


def test_formula_row_with_empty_inputs_takes_db_total_but_keeps_formula_otherwise():
    wb = Workbook()
    ws = wb.active
    ws["A4"], ws["C4"], ws["D4"] = "项目", "期末余额", "期初余额"
    ws["A5"], ws["B5"], ws["C5"], ws["D5"] = "实收资本", "七、58", "=SUM(C6:C7)", "=SUM(D6:D7)"
    ws["A6"], ws["C6"], ws["D6"] = "国家资本", 0, None
    ws["A7"], ws["C7"], ws["D7"] = "外商资本", None, 5
    best = {("BS", "实收资本", ""): {"row_code": "BS-081", "current_period_amount": Decimal("150"),
                                     "prior_period_amount": Decimal("150")}}
    assert fill_formula_rows_without_inputs(ws, best, "BS") == 1
    assert ws["C5"].value == pytest.approx(150)  # 明细全空 → 库内总额（附注编号格不被当成行名）
    assert ws["D5"].value == "=SUM(D6:D7)"  # 明细有数 → 保留模板公式


def test_consolidated_only_formula_rows_blank_in_standalone():
    wb = Workbook()
    ws = wb.active
    ws["A4"], ws["C4"], ws["D4"] = "项目", "本期金额", "上期金额"
    ws["A5"], ws["C5"] = "五、净利润", 30
    ws["A6"], ws["C6"] = "1.归属于母公司所有者的净利润", None
    ws["A7"], ws["C7"], ws["D7"] = "  *2.少数股东损益", "=C5-C6", "=D5-D6"
    ws["A8"], ws["C8"] = "  *少数股东权益", 7  # 库内填了数的格不动
    assert blank_consolidated_only_formulas(ws) == 2
    assert ws["C7"].value is None and ws["D7"].value is None
    assert ws["C8"].value == 7

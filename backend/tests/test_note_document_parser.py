"""Phase 3 测试：DOCX / PDF / 纯文本解析器。"""

from __future__ import annotations

import io
import zipfile

import pytest

from app.services.note_document_parser import (
    ParseError,
    parse_docx,
    parse_pdf,
    parse_text_source,
)


# ── DOCX 辅助 ─────────────────────────────────────────────────────────────

def _minimal_docx(paragraphs: list[tuple[str, str | None]]) -> bytes:
    """生成最小 DOCX：paragraphs = [(text, style_name | None), ...]"""
    from docx import Document
    from docx.enum.text import WD_ALIGN_PARAGRAPH

    doc = Document()
    for text, style in paragraphs:
        if style and ("Heading" in style or "标题" in style):
            doc.add_heading(text, level=int(style.split()[-1]) if style.split()[-1].isdigit() else 1)
        else:
            doc.add_paragraph(text)
    buf = io.BytesIO()
    doc.save(buf)
    return buf.getvalue()


class TestParseDocx:
    def test_basic_heading_and_body(self):
        data = _minimal_docx([
            ("一、公司基本情况", "Heading 1"),
            ("本公司成立于2020年。", None),
            ("二、重要会计政策", "Heading 1"),
            ("会计期间为1月至12月。", None),
        ])
        result = parse_docx(data)
        assert len(result.sections) >= 2
        assert result.sections[0].title == "一、公司基本情况"
        assert "本公司" in result.sections[0].html
        assert result.file_hash

    def test_no_heading_creates_untitled(self):
        data = _minimal_docx([("普通段落没有标题", None)])
        result = parse_docx(data)
        assert len(result.sections) >= 1
        assert result.sections[0].confidence == "manual_required"

    def test_empty_docx_raises(self):
        from docx import Document
        doc = Document()
        buf = io.BytesIO()
        doc.save(buf)
        with pytest.raises(ParseError, match="NO_PARSEABLE_TEXT"):
            parse_docx(buf.getvalue())

    def test_corrupt_file_raises(self):
        with pytest.raises(ParseError, match="UNSUPPORTED_FILE"):
            parse_docx(b"this is not a docx file")

    def test_zip_bomb_raises(self):
        buf = io.BytesIO()
        with zipfile.ZipFile(buf, "w") as zf:
            for i in range(2500):
                zf.writestr(f"entry_{i}.txt", "x" * 100)
        with pytest.raises(ParseError, match="DOCX_ZIP_LIMIT"):
            parse_docx(buf.getvalue())

    def test_html_is_sanitized(self):
        data = _minimal_docx([
            ("一、标题", "Heading 1"),
            ("正文内容", None),
        ])
        result = parse_docx(data)
        for s in result.sections:
            assert "<script>" not in s.html


# ── PDF 辅助 ──────────────────────────────────────────────────────────────

def _minimal_pdf(text: str = "一、公司基本情况\n本公司成立于2020年。") -> bytes:
    """用 PyMuPDF 生成最小 PDF。"""
    import pymupdf
    doc = pymupdf.open()
    page = doc.new_page()
    page.insert_text((72, 72), text, fontsize=14, fontname="china-s")
    data = doc.tobytes()
    doc.close()
    return data


class TestParsePdf:
    def test_basic_text_pdf(self):
        data = _minimal_pdf()
        result = parse_pdf(data)
        # 至少有一个章节或降级
        assert result.sections or result.needs_ocr
        assert result.file_hash

    def test_encrypted_pdf_raises(self):
        import pymupdf
        doc = pymupdf.open()
        page = doc.new_page()
        page.insert_text((72, 72), "内容")
        encrypted = doc.tobytes(encryption=pymupdf.PDF_ENCRYPT_AES_256, user_pw="secret")
        doc.close()
        with pytest.raises(ParseError, match="已加密"):
            parse_pdf(encrypted)

    def test_empty_pdf_needs_ocr(self):
        """空白 PDF（无文本 span）应标 needs_ocr。"""
        import pymupdf
        doc = pymupdf.open()
        doc.new_page()  # 空白页
        data = doc.tobytes()
        doc.close()
        result = parse_pdf(data)
        assert result.needs_ocr is True

    def test_corrupt_pdf_raises(self):
        with pytest.raises(ParseError, match="UNSUPPORTED_FILE"):
            parse_pdf(b"not a pdf")


# ── 文本来源 ─────────────────────────────────────────────────────────────

class TestParseTextSource:
    def test_markdown_headings(self):
        text = "# 一、概述\n公司介绍\n## 二、政策\n会计政策说明"
        result = parse_text_source(text)
        assert len(result.sections) == 2
        assert result.sections[0].title == "一、概述"

    def test_html_headings(self):
        text = "<h1>一、概述</h1><p>公司介绍</p><h2>二、政策</h2><p>说明</p>"
        result = parse_text_source(text)
        assert len(result.sections) == 2

    def test_cn_headings(self):
        text = "一、公司基本情况\n公司介绍\n二、重要会计政策\n政策说明"
        result = parse_text_source(text)
        assert len(result.sections) == 2

    def test_single_paragraph_fallback(self):
        text = "这是一段没有任何标题结构的纯文本内容用于测试降级路径"
        result = parse_text_source(text, doc_name="测试文档.md")
        assert len(result.sections) >= 1
        # 无标题时降级为单段或 cn 候选；重点是不崩溃且有内容

    def test_empty_text(self):
        result = parse_text_source("")
        assert result.sections == []
        assert any("为空" in w for w in result.warnings)

    def test_html_sanitized(self):
        text = "# 标题\n<script>alert(1)</script>正文"
        result = parse_text_source(text)
        for s in result.sections:
            # 消毒后 <script> 标签不应以原始形式存在；转义为 &lt;script&gt; 是安全的
            assert "<script>" not in s.html

"""底稿表头填充守卫（``wp_header_service``，生成底稿时写编制单位 / 期间 / 截止日）。

2026-10-26 修「按编码生成时，遇到有合并单元格的模板，表头会静默填不上」：
- 根因：openpyxl 里合并区域除左上角外都是只读的 ``MergedCell``；原先直接写「标签右边一格」、没找到标签就在活动页
  固定格写标准布局，一碰合并格 ``AttributeError``，调用方只记 warning、回执照报成功（按生成链路实测 J2A / K11-4 /
  B10-2 / L1-1 / J2-5 / J2-6 / A10 / A12 / B13-2 全是这样）；
- 同时修掉：不带冒号的表格列头（程序表「索引号」）右边被写底稿编码；标准布局兜底盖掉模板第 1~5 行的内容；
  表头页（底稿目录 / 表头（请先填写） / 目录）的右边是目录表格，只能写在标签冒号后。
"""
from __future__ import annotations

import hashlib
import re
import uuid
from pathlib import Path
from types import SimpleNamespace
from typing import Any

import openpyxl
import pytest
from openpyxl.cell.cell import MergedCell

from app.services import wp_header_service as hdr
from app.services import wp_template_init_service as init

CLIENT = "某某测试有限公司"
PERIOD = "2025年01月01日至2025年12月31日"
CUTOFF = "2025年12月31日"
VALUES = {"client": CLIENT, "period": PERIOD, "cutoff": CUTOFF, "wp_name": "名称X", "code": "J2A",
          "xref": "", "preparer": "", "reviewer": ""}
HUB_A1 = "底稿开列时，可先在以下单元格填写有关内容，以便其他底稿引用"


def _book(cells: dict[str, Any], merges: tuple[str, ...] = (), title: str = "审定表"):
    wb = openpyxl.Workbook()
    ws = wb.active
    ws.title = title
    for coord, value in cells.items():
        ws[coord] = value
    for rng in merges:
        ws.merge_cells(rng)
    return wb


def _add_sheet(wb, title: str, cells: dict[str, Any]):
    ws = wb.create_sheet(title)
    for coord, value in cells.items():
        ws[coord] = value
    return ws


def _written(filled: list[dict[str, str]]) -> set[tuple[str, str, str]]:
    return {(f["sheet"], f["cell"], f["field"]) for f in filled}


# ── 合并格 ──────────────────────────────────────────────────────────────────────


def test_merged_label_value_goes_to_first_column_after_label_range():
    """「被审计单位：」占 A3:B3 ⇒ 右边一格 B3 是只读合并格（原先 AttributeError），值写 C3。"""
    wb = _book({"A3": "被审计单位：", "A4": "会计期间："}, merges=("A3:B3", "A4:B4"))
    _ext, filled = hdr._fill_workbook_header(wb, VALUES)
    ws = wb.active
    assert ws["C3"].value == CLIENT and ws["C4"].value == PERIOD
    assert _written(filled) == {("审定表", "C3", "client"), ("审定表", "C4", "period")}


def test_merged_value_region_is_written_at_its_anchor():
    """值区 B3:E3 合并 ⇒ 写左上角 B3（显示在整块里）。"""
    wb = _book({"A3": "编制单位：", "A4": "截止日："}, merges=("B3:E3", "B4:E4"))
    hdr._fill_workbook_header(wb, VALUES)
    assert wb.active["B3"].value == CLIENT and wb.active["B4"].value == CUTOFF


def test_vertical_value_region_shared_by_two_labels_is_not_overwritten():
    """值区 B3:B4 竖向合并、两行各一个标签：B4 是只读合并格（原先 AttributeError）⇒ 第一个标签写左上角 B3，
    第二个落到同一格时改写进自己冒号后，不覆盖前一个。"""
    wb = _book({"A3": "被审计单位：", "A4": "截止日："}, merges=("B3:B4",))
    _ext, filled = hdr._fill_workbook_header(wb, VALUES)
    ws = wb.active
    assert ws["B3"].value == CLIENT and ws["A4"].value == f"截止日：{CUTOFF}"
    assert _written(filled) == {("审定表", "B3", "client"), ("审定表", "A4", "cutoff")}


def test_value_region_merged_with_label_neighbour_label_range_ending_inside():
    """标签 A3:C3 合并、值区 D3:F3 合并 ⇒ 写 D3；不碰 B3 / C3 / E3（都是只读合并格）。"""
    wb = _book({"A3": "被审计单位名称："}, merges=("A3:C3", "D3:F3"))
    hdr._fill_workbook_header(wb, VALUES)
    ws = wb.active
    assert ws["D3"].value == CLIENT
    assert isinstance(ws["B3"], MergedCell) and isinstance(ws["E3"], MergedCell)


# ── 标签认法 ───────────────────────────────────────────────────────────────────


def test_column_headers_without_colon_are_not_labels():
    """程序表表头行「索引号」「复核人」「工作底稿索引」不带冒号 ⇒ 不写（原先把底稿编码写进「索引号」列头右边）。"""
    wb = _book({"G5": "索引号", "H5": None, "I5": "复核人", "A6": "工作底稿索引", "A3": "被审计单位："})
    _ext, filled = hdr._fill_workbook_header(wb, VALUES)
    ws = wb.active
    assert ws["H5"].value is None and ws["J5"].value is None and ws["B6"].value is None
    assert _written(filled) == {("审定表", "B3", "client")}


def test_label_name_must_match_exactly():
    """按全名认标签：以关键词结尾（「上年审计期间：」「工作底稿索引号：」，原先按结尾匹配会写）或开头的都不算。"""
    wb = _book({"A3": "被函证单位名称：", "A4": "上年审计期间：", "A5": "工作底稿索引号：", "A6": "编制人员数：",
                "A7": "截止日期："})
    _ext, filled = hdr._fill_workbook_header(wb, VALUES)
    ws = wb.active
    assert [ws[f"B{r}"].value for r in range(3, 7)] == [None, None, None, None]
    assert _written(filled) == {("审定表", "B7", "cutoff")}


def test_inline_examples_replaced_only_for_client_period_cutoff():
    """冒号后是模板示例：客户名称 / 期间 / 截止日换成本项目的值；索引号（模板自己的）与编制人（生成时没有值）不动。"""
    wb = _book({"A2": "被审计单位：ABC公司", "A3": "会计期间：201X年度", "A4": "截止日：202X年12月31日",
                "A5": "索引号：A10", "A6": "编制人：张三", "A7": "客户名称：  "})
    hdr._fill_workbook_header(wb, VALUES)
    ws = wb.active
    assert ws["A2"].value == f"被审计单位：{CLIENT}"
    assert ws["A3"].value == f"会计期间：{PERIOD}"
    assert ws["A4"].value == f"截止日：{CUTOFF}"
    assert ws["A5"].value == "索引号：A10" and ws["A6"].value == "编制人：张三"
    assert ws["B7"].value == CLIENT, "冒号后只有空白 ⇒ 当作空标签，写右边"


def test_formula_or_template_value_on_the_right_is_kept():
    """右边是公式（引用表头页）或模板自带的值（索引号 A10） ⇒ 不动。"""
    wb = _book({"A3": "编制单位：", "B3": "=底稿目录!A2", "A4": "索引号：", "B4": "A10"})
    _add_sheet(wb, "底稿目录", {"A1": HUB_A1, "A2": "被审计单位："})
    hdr._fill_workbook_header(wb, {**VALUES, "code": "E1-1"})
    ws = wb["审定表"]
    assert ws["B3"].value == "=底稿目录!A2" and ws["B4"].value == "A10"


def test_label_followed_by_another_label_is_filled_after_its_colon():
    """「客户名称：」右边紧挨「编制人：」 ⇒ 写进标签格冒号后，右边的标签不动。"""
    wb = _book({"A2": "客户名称：", "B2": "编制人：", "A3": "会计时点：", "A4": "截止日/会计期间："})
    hdr._fill_workbook_header(wb, VALUES)
    ws = wb.active
    assert ws["A2"].value == f"客户名称：{CLIENT}" and ws["B2"].value == "编制人："
    assert ws["B4"].value == PERIOD, "「截止日/会计期间」整名登记为期间"
    assert ws["B3"].value is None, "「会计时点」未登记，不猜"


# ── 表头页与跨页引用 ─────────────────────────────────────────────────────────────


def _hub_book(hub_title: str, hub_a1: str | None):
    """活动页「审定表」的表头格全是引用表头页的公式；表头页右边是目录表格（序号 / 内容）。"""
    wb = _book({"A1": "致同会计师事务所", "A3": f"='{hub_title}'!A2", "C3": f"='{hub_title}'!A4",
                "A4": f"='{hub_title}'!A3", "G3": "索引号：", "H3": "E1-1"})
    _add_sheet(wb, hub_title, {"A1": hub_a1, "A2": "被审计单位：", "A3": "截止日：202X年12月31日", "A4": "编制人：",
                               "D2": "序号", "E2": "内  容", "D3": 1, "E3": "审定表E1-1"})
    return wb


@pytest.mark.parametrize(("hub_title", "hub_a1"), [("底稿目录", None), ("表头（请先填写）", HUB_A1), ("目录", HUB_A1)])
def test_hub_sheet_is_filled_after_colon_and_directory_table_untouched(hub_title, hub_a1):
    wb = _hub_book(hub_title, hub_a1)
    _ext, filled = hdr._fill_workbook_header(wb, {**VALUES, "code": "E1-1"})
    hub, ws = wb[hub_title], wb["审定表"]
    assert hub["A2"].value == f"被审计单位：{CLIENT}" and hub["A3"].value == f"截止日：{CUTOFF}"
    assert hub["A4"].value == "编制人："
    assert (hub["B2"].value, hub["D2"].value, hub["E2"].value, hub["E3"].value) == (None, "序号", "内  容", "审定表E1-1")
    assert ws["A3"].value == f"='{hub_title}'!A2", "活动页的引用公式不动（显示的就是表头页填好的值）"
    assert ws["H3"].value == "E1-1"
    assert _written(filled) == {(hub_title, "A2", "client"), (hub_title, "A3", "cutoff")}


def test_hub_recognised_by_a1_note_even_when_it_is_the_active_sheet():
    """活动页就是表头页（K11-4 / E1 / D2 生成后活动页是「底稿目录」）⇒ 同样只写冒号后。"""
    wb = _hub_book("表头", HUB_A1)
    wb.active = wb.sheetnames.index("表头")
    hdr._fill_workbook_header(wb, {**VALUES, "code": "E1-1"})
    hub = wb["表头"]
    assert hub["A2"].value == f"被审计单位：{CLIENT}" and hub["B2"].value is None and hub["D2"].value == "序号"


def test_two_hop_reference_reaches_the_label_cell():
    """活动页 → 纵向分析 → 横向分析的标签格（A1-13 / B13-2 这类多页分析底稿）：写最终那一格的冒号后，中间的公式不动。"""
    wb = _book({"A2": "='纵向分析'!A2", "B2": "='纵向分析'!E2"}, title="审定表")
    _add_sheet(wb, "纵向分析", {"A2": "='横向分析'!A2", "E2": "='横向分析'!D2"})
    _add_sheet(wb, "横向分析", {"A2": "客户名称：", "D2": "编制人：", "A3": "会计时点："})
    _ext, filled = hdr._fill_workbook_header(wb, VALUES)
    assert wb["横向分析"]["A2"].value == f"客户名称：{CLIENT}"
    assert wb["横向分析"]["B2"].value is None, "被引用的格只写冒号后，不写右边"
    assert wb["纵向分析"]["A2"].value == "='横向分析'!A2"
    assert _written(filled) == {("横向分析", "A2", "client")}


def test_reference_to_missing_sheet_or_compound_formula_is_ignored():
    wb = _book({"A2": "='不存在'!A2", "A3": "=底稿目录!A2&\"\"", "A4": "被审计单位："})
    _ext, filled = hdr._fill_workbook_header(wb, VALUES)
    assert _written(filled) == {("审定表", "B4", "client")}


def test_other_sheets_get_client_period_cutoff_after_colon_only():
    """非活动页（既不是表头页、也没被活动页引用）：编制单位 / 期间 / 截止日写进冒号后、示例换成本项目值；
    索引号 / 底稿名称各页不同不动，右边一格不写（spec post-push-residual-defect-closure R3.4：
    活动页是「编制说明」的 A3-7、各页自带表头的 E0 原先其余各页仍显示「截止日：202X年12月31日」）。"""
    wb = _book({"A1": "说明"}, title="编制说明")
    _add_sheet(wb, "明细表", {"A3": "客户名称：", "A4": "截止日：202X年12月31日", "H3": "索引号：",
                              "A5": "底稿名称：", "D3": "编制人："})
    _ext, filled = hdr._fill_workbook_header(wb, VALUES)
    ws = wb["明细表"]
    assert ws["A3"].value == f"客户名称：{CLIENT}" and ws["A4"].value == f"截止日：{CUTOFF}"
    assert (ws["B3"].value, ws["H3"].value, ws["I3"].value, ws["A5"].value, ws["D3"].value) == (
        None, "索引号：", None, "底稿名称：", "编制人：")
    assert _written(filled) == {("明细表", "A3", "client"), ("明细表", "A4", "cutoff")}


# ── fill_workpaper_header：保存与回报 ──────────────────────────────────────────────


class _Result:
    def __init__(self, obj):
        self._obj = obj

    def scalar_one_or_none(self):
        return self._obj


class _FakeDB:
    """只回答「按 id 取项目」这一条查询（fill_workpaper_header 唯一的数据库访问）。"""

    def __init__(self, project):
        self.project = project

    async def execute(self, *_args, **_kwargs):
        return _Result(self.project)


def _db(**overrides) -> _FakeDB:
    project = dict(id=uuid.uuid4(), name="【测试】表头填充", client_name=CLIENT, audit_year=2025,
                   audit_period_start=None, audit_period_end=None, wizard_state=None)
    project.update(overrides)
    return _FakeDB(SimpleNamespace(**project))


async def _fill(path: Path, *, code: str = "J2A", db: _FakeDB | None = None, **kwargs) -> dict:
    return await hdr.fill_workpaper_header(
        db=db or _db(), project_id=uuid.uuid4(), wp_id=uuid.uuid4(), file_path=str(path),
        wp_code=code, wp_name="名称X", cycle=code[0], **kwargs,
    )


def _sha(path: Path) -> str:
    return hashlib.sha256(path.read_bytes()).hexdigest()


async def test_merged_template_header_is_saved(tmp_path):
    path = tmp_path / "t.xlsx"
    _book({"A3": "被审计单位：", "A4": "会计期间："}, merges=("A3:B3", "A4:B4")).save(path)
    res = await _fill(path)
    assert res["status"] == "ok" and res["standard_layout"] is False
    wb = openpyxl.load_workbook(path)
    assert (wb.active["C3"].value, wb.active["C4"].value) == (CLIENT, PERIOD)


async def test_template_without_header_labels_is_left_byte_identical(tmp_path):
    """模板里没有可填的表头（说明页 / 指引页）⇒ 成功、什么都不写、也不重存（openpyxl 重存会丢数据验证扩展等）。
    原先这里写标准布局，盖掉模板第 1~5 行（A2-1「说明」页、A3-5「目录」页）。"""
    path = tmp_path / "t.xlsx"
    _book({"A1": "说明", "A2": "本表用于汇总各单体报表", "B3": "序号", "C3": "内容"}, title="说明").save(path)
    before = _sha(path)
    res = await _fill(path, code="A2-1")
    assert res == {"status": "ok", "message": None, "filled": [], "standard_layout": False, "filled_by_search": False}
    assert _sha(path) == before


async def test_blank_workbook_gets_standard_layout(tmp_path):
    """找不到模板时建的空白册（from_template=False）⇒ 按标准布局写表头。"""
    path = tmp_path / "t.xlsx"
    openpyxl.Workbook().save(path)
    res = await _fill(path, code="ZZ9", from_template=False)
    assert res["status"] == "ok" and res["standard_layout"] is True
    ws = openpyxl.load_workbook(path).active
    assert ws["A1"].value == "致同会计师事务所（特殊普通合伙）" and ws["B2"].value == CLIENT and ws["F3"].value == "ZZ9"


async def test_custom_workpaper_always_gets_standard_layout(tmp_path):
    path = tmp_path / "t.xlsx"
    _book({"A3": "被审计单位："}).save(path)
    res = await _fill(path, code="ZZ1", is_custom=True)
    assert res["standard_layout"] is True
    assert openpyxl.load_workbook(path).active["B2"].value == CLIENT


async def test_external_link_header_is_reported_not_silent(tmp_path):
    """表头格全是外部链接（=[4]底稿目录!$A$2：D2-5 / D4-12 这类从别的文件拆出来的模板）⇒ partial + 原因，文件不动。"""
    path = tmp_path / "t.xlsx"
    _book({"A3": "=[4]底稿目录!$A$2", "C3": "=[4]底稿目录!$A$4", "A4": "=[4]底稿目录!$A$3",
           "G3": "索引号：", "H3": "D2-5"}).save(path)
    before = _sha(path)
    res = await _fill(path, code="D2-5")
    assert res["status"] == "partial" and "外部链接" in res["message"] and res["filled"] == []
    assert _sha(path) == before


async def test_external_link_not_reported_when_client_is_shown_elsewhere_on_active_sheet(tmp_path):
    path = tmp_path / "t.xlsx"
    _book({"A3": "被审计单位：", "C3": "=[4]底稿目录!$A$4"}).save(path)
    res = await _fill(path)
    assert res["status"] == "ok" and res["message"] is None
    assert openpyxl.load_workbook(path).active["B3"].value == CLIENT


@pytest.mark.parametrize(("make", "message"), [
    (lambda p: p, "底稿文件不存在"),
    (lambda p: (p.write_bytes(b"not a workbook"), p)[1], "打开底稿文件失败"),
])
async def test_errors_are_reported_without_paths_or_exception_text(tmp_path, make, message):
    """原因会进生成回执给用户看：只写原因，不带服务器路径与异常原文（那些记日志）。"""
    path = make(tmp_path / "t.xlsx")
    res = await _fill(path)
    assert res == {"status": "error", "message": message}
    assert str(tmp_path) not in res["message"]


# ── 表头外部链接引用归类（spec post-push-residual-defect-closure R3）────────────────────────────────


def _with_cached(path: Path, cached: dict[tuple[str, str], str]) -> None:
    """给公式格写缓存值（模拟模板里外部链接的缓存显示值；openpyxl 存盘不写缓存）。"""
    from app.services.xlsx_cell_patch import CellPatch, patch_cells

    patch_cells(path, [CellPatch(s, c, cached=v) for (s, c), v in cached.items()])


def _formula_and_shown(path: Path, sheet: str, cell: str) -> tuple[Any, Any]:
    wb, wbv = openpyxl.load_workbook(path), openpyxl.load_workbook(path, data_only=True)
    try:
        return wb[sheet][cell].value, wbv[sheet][cell].value
    finally:
        wb.close()
        wbv.close()


async def test_external_ref_to_same_label_in_this_book_becomes_internal_ref(tmp_path):
    """缓存是「被审计单位：ABC公司」、本册底稿目录 A2 是同一标签 ⇒ 改成本册引用，表头页那一格填值，
    引用格的缓存值随之更新（按缓存值读的 HTML 渲染与重算后一致）。"""
    path = tmp_path / "t.xlsx"
    wb = _book({"A3": "=[4]底稿目录!A2", "A4": "='[4]底稿目录'!$A$3"})
    _add_sheet(wb, "底稿目录", {"A2": "被审计单位：", "A3": "截止日：202X年12月31日"})
    wb.save(path)
    _with_cached(path, {("审定表", "A3"): "被审计单位：ABC公司", ("审定表", "A4"): "截止日：202X年12月31日"})
    res = await _fill(path, code="E1-28")
    assert res["status"] == "ok", res
    assert _formula_and_shown(path, "审定表", "A3") == ("=底稿目录!A2", f"被审计单位：{CLIENT}")
    assert _formula_and_shown(path, "审定表", "A4") == ("='底稿目录'!$A$3", f"截止日：{CUTOFF}")
    assert _formula_and_shown(path, "底稿目录", "A2")[0] == f"被审计单位：{CLIENT}"


async def test_external_label_without_matching_cell_becomes_static_label(tmp_path):
    """本册没有被引用的页（D2-5 / D4-12 这类拆出来的模板）⇒ 改成静态标签、示例去掉，能填的项填本项目值；
    编制人 / 日期这类只留标签。"""
    path = tmp_path / "t.xlsx"
    _book({"A3": "=[1]底稿目录!$A$2", "A4": "=[1]底稿目录!$A$3", "H3": "=[1]底稿目录!$A$4",
           "H4": "=[1]底稿目录!$A$5"}).save(path)
    _with_cached(path, {("审定表", "A3"): "被审计单位：", ("审定表", "A4"): "会计期间：202X年度",
                        ("审定表", "H3"): "编制人：张XX", ("审定表", "H4"): "编制日期：2025-4-1"})
    res = await _fill(path, code="D4-5")
    assert res["status"] == "ok", res
    ws = openpyxl.load_workbook(path).active
    assert (ws["A3"].value, ws["A4"].value) == (f"被审计单位：{CLIENT}", f"会计期间：{PERIOD}")
    assert (ws["H3"].value, ws["H4"].value) == ("编制人：", "编制日期：")


async def test_external_index_code_becomes_static_text_and_data_refs_untouched(tmp_path):
    """缓存是本页索引号（F2-18）⇒ 静态文字；缓存是金额 / 名称（数据）⇒ 公式与缓存都不动。"""
    path = tmp_path / "t.xlsx"
    _book({"A3": "被审计单位：", "I3": "=[3]底稿目录!F2", "F8": "='[13]可收回金额测试D27-6'!N8",
           "B9": "=[2]银行!B4"}).save(path)
    _with_cached(path, {("审定表", "I3"): "F2-18", ("审定表", "F8"): "1234.5", ("审定表", "B9"): "XX银行"})
    res = await _fill(path, code="F2")
    assert res["status"] == "ok", res
    assert _formula_and_shown(path, "审定表", "I3") == ("F2-18", "F2-18")
    assert _formula_and_shown(path, "审定表", "F8")[0] == "='[13]可收回金额测试D27-6'!N8"
    assert _formula_and_shown(path, "审定表", "B9") == ("=[2]银行!B4", "XX银行")


async def test_long_reference_chain_and_same_sheet_refs_get_new_cached_values(tmp_path):
    """B13-4 的表头经 4 页转引才到底稿目录、右侧再 ``=A2`` 引一次：按缓存值读（HTML 渲染）都要显示本项目值
    （原先只跟 3 层、不认同页引用 ⇒ 仍显示「被审计单位：ABC公司」）。"""
    path = tmp_path / "t.xlsx"
    wb = _book({"A2": "='第4页'!A2", "J2": "=A2"}, title="第5页")
    for i in (4, 3, 2):
        _add_sheet(wb, f"第{i}页", {"A2": f"='第{i - 1}页'!A2"})
    _add_sheet(wb, "第1页", {"A2": "=底稿目录!A2"})
    _add_sheet(wb, "底稿目录", {"A2": "被审计单位：ABC公司"})
    wb.save(path)
    _with_cached(path, {(f"第{i}页", "A2"): "被审计单位：ABC公司" for i in range(1, 6)}
                 | {("第5页", "J2"): "被审计单位：ABC公司"})
    res = await _fill(path, code="B13-4")
    assert res["status"] == "ok", res
    for title, cell in [(f"第{i}页", "A2") for i in range(1, 6)] + [("第5页", "J2")]:
        assert _formula_and_shown(path, title, cell)[1] == f"被审计单位：{CLIENT}", (title, cell)


async def test_unresolvable_external_header_is_still_reported(tmp_path):
    """缓存为空、本册也找不到对应格的外部引用（链接已断）⇒ 不动，客户名不可见时照旧回报 partial。"""
    path = tmp_path / "t.xlsx"
    _book({"A3": "=[4]表头!A2", "G3": "索引号：", "H3": "D2-5"}).save(path)
    res = await _fill(path, code="D2-5")
    assert res["status"] == "partial" and "外部链接" in res["message"]
    assert _formula_and_shown(path, "审定表", "A3")[0] == "=[4]表头!A2"


async def test_generated_xlsm_header_is_filled(tmp_path):
    """含宏底稿（.xlsm）同样填表头（只改目标格，宏原样保留；原先因怕 openpyxl 丢宏直接跳过）。"""
    path = tmp_path / "t.xlsm"
    _book({"A3": "被审计单位："}).save(path)
    issue = await hdr.fill_generated_workpaper_header(
        _db(), project_id=uuid.uuid4(), wp_id=uuid.uuid4(), path=path, wp_code="B23-1", wp_name="名称X",
        cycle="B", from_template=True,
    )
    assert issue is None
    assert openpyxl.load_workbook(path).active["B3"].value == CLIENT


async def test_period_and_cutoff_come_from_project_dates(tmp_path):
    from datetime import date

    path = tmp_path / "t.xlsx"
    _book({"A3": "会计期间：", "A4": "截止日："}).save(path)
    await _fill(path, db=_db(audit_period_start=date(2025, 4, 1), audit_period_end=date(2026, 3, 31)))
    ws = openpyxl.load_workbook(path).active
    assert ws["B3"].value == "2025年04月01日至2026年03月31日" and ws["B4"].value == "2026年03月31日"


# ── 真实模板（按生成链路建底稿，存储目录指到临时目录） ─────────────────────────────────────

#: 原先在合并格上抛 AttributeError、表头一格没填的编码（按生成链路逐个实测）+ 表头页 / A1 说明表头页 / 多文件合并的典型
_MERGED_CRASH_CODES = ("J2A", "K11-4", "B10-2", "L1-1", "J2-5", "J2-6", "A10", "A12", "B13-2")
_OTHER_LAYOUT_CODES = ("C21", "A5-3", "S34-16", "E1")


def _snapshot(wb) -> dict[tuple[str, str], Any]:
    return {(s.title, c.coordinate): c.value for s in wb.worksheets
            for row in s.iter_rows(min_row=1, max_row=hdr._HEADER_ROWS, max_col=hdr._HEADER_COLS) for c in row
            if not isinstance(c, MergedCell) and c.value not in (None, "")}


def _shown(wb, value: Any, depth: int = 0) -> Any:
    hop = hdr._follow_reference(wb, value) if depth < 4 else None
    if hop is None:
        return value
    cell = hop[0][hop[1]]
    return _shown(wb, None if isinstance(cell, MergedCell) else cell.value, depth + 1)


def _generate(tmp_path: Path, monkeypatch, code: str) -> Path:
    monkeypatch.setattr(init, "STORAGE_DIR", tmp_path / "storage")
    pid = uuid.uuid4()
    (tmp_path / "storage" / str(pid)).mkdir(parents=True)
    path = init.init_workpaper_from_template(pid, uuid.uuid4(), code)
    if path is None:
        pytest.fail(f"{code} 找不到模板（模板目录变了？）")
    assert path.suffix == ".xlsx", f"{code} 生成的不是 xlsx：{path.name}"
    return path


@pytest.mark.parametrize("code", _MERGED_CRASH_CODES + _OTHER_LAYOUT_CODES)
async def test_real_template_header_filled_without_touching_other_content(tmp_path, monkeypatch, code):
    path = _generate(tmp_path, monkeypatch, code)
    wb = openpyxl.load_workbook(path)
    before, sheets = _snapshot(wb), wb.sheetnames
    wb.close()

    res = await _fill(path, code=code)
    assert res["status"] == "ok", res
    assert res["filled"], f"{code} 一格都没填"

    wb = openpyxl.load_workbook(path)
    assert wb.sheetnames == sheets
    after = _snapshot(wb)
    for key, old in before.items():
        new = after.get(key)
        if new == old or _external_rewrite_ok(old, new):
            continue
        prefix = re.match(r"^(.*?[：:])", str(old))
        assert prefix and isinstance(new, str) and new.startswith(prefix.group(1)), (
            f"{code} {key} 原有内容被改：{old!r} → {new!r}（只许在标签冒号后改写）")
    ws = wb.active
    shown = [_shown(wb, c.value) for row in ws.iter_rows(min_row=1, max_row=hdr._HEADER_ROWS, max_col=hdr._HEADER_COLS)
             for c in row if not isinstance(c, MergedCell)]
    assert any(isinstance(v, str) and CLIENT in v for v in shown), f"{code} 活动页「{ws.title}」表头看不到编制单位"


def _external_rewrite_ok(old: Any, new: Any) -> bool:
    """表头外部链接单格引用只许改成本册引用（去掉 ``[N]``，页签与地址不变）或静态文字（标签 / 索引号）。"""
    if not (isinstance(old, str) and hdr._EXTERNAL_SINGLE_REF_RE.match(old)):
        return False
    if isinstance(new, str) and new.startswith("="):
        return new == re.sub(r"\[\d+\]", "", old, count=1)
    return isinstance(new, str) and not new.startswith("=")


#: 表头显示模板示例的痕迹（外部链接缓存值）：「会计期间：202X年度」「被审计单位：ABC公司」「编制人：张XX」
_EXAMPLE_TRACE = re.compile(r"202X|201X|ABC公司|张XX|李XX|2025-4-1|2024-4-1")


def _displayed_header(path: Path) -> dict[tuple[str, str], Any]:
    """按缓存值读（与 HTML 渲染同口径）：各页表头区非空格的显示值。"""
    wb = openpyxl.load_workbook(path, data_only=True)
    try:
        return {(s.title, c.coordinate): c.value for s in wb.worksheets
                for row in s.iter_rows(min_row=1, max_row=hdr._HEADER_ROWS, max_col=hdr._HEADER_COLS) for c in row
                if not isinstance(c, MergedCell) and c.value not in (None, "")}
    finally:
        wb.close()


def _external_data_refs(path: Path) -> dict[tuple[str, str], tuple[str, Any]]:
    """表头区缓存是数据（金额 / 银行名称，不是表头标签也不是索引号）的外部链接格：(公式, 缓存值)。"""
    wb, wbv = openpyxl.load_workbook(path), openpyxl.load_workbook(path, data_only=True)
    try:
        out = {}
        for s in wb.worksheets:
            for row in s.iter_rows(min_row=1, max_row=hdr._HEADER_ROWS, max_col=hdr._HEADER_COLS):
                for c in row:
                    if isinstance(c, MergedCell) or not isinstance(c.value, str) or not hdr._EXTERNAL_REF_RE.match(c.value):
                        continue
                    cached = wbv[s.title][c.coordinate].value
                    if cached is None or hdr._header_label(cached):
                        continue
                    if isinstance(cached, str) and hdr._WP_CODE_TEXT_RE.match(cached.strip()):
                        continue
                    out[(s.title, c.coordinate)] = (c.value, cached)
        return out
    finally:
        wb.close()
        wbv.close()


def _zip_parts(path: Path) -> dict[str, bytes]:
    import zipfile

    with zipfile.ZipFile(path) as zf:
        return {n: zf.read(n) for n in zf.namelist()}


@pytest.mark.parametrize("code", ["H4", "C2", "B23-1"])
async def test_real_template_header_fill_keeps_data_validations_and_macros(tmp_path, monkeypatch, code):
    """按编码生成填表头只改目标格（spec post-push-residual-defect-closure R4）：x14 数据验证条数、宏、
    没改到的部件逐字节不变（原先 openpyxl 整册重存丢 x14 数据验证与全部公式缓存值，.xlsm 干脆不填）。"""
    monkeypatch.setattr(init, "STORAGE_DIR", tmp_path / "storage")
    pid = uuid.uuid4()
    (tmp_path / "storage" / str(pid)).mkdir(parents=True)
    path = init.init_workpaper_from_template(pid, uuid.uuid4(), code)
    assert path is not None
    before = _zip_parts(path)
    dv = sum(v.count(b"<x14:dataValidation ") for n, v in before.items() if n.startswith("xl/worksheets/"))
    assert dv > 0, f"{code} 应含 x14 数据验证（模板变了？）"
    issue = await hdr.fill_generated_workpaper_header(
        _db(), project_id=pid, wp_id=uuid.uuid4(), path=Path(path), wp_code=code, wp_name="名称X",
        cycle=code[0], from_template=True,
    )
    assert issue is None
    after = _zip_parts(path)
    assert sum(v.count(b"<x14:dataValidation ") for n, v in after.items() if n.startswith("xl/worksheets/")) == dv
    changed = [n for n in before if before[n] != after.get(n)]
    assert changed and all(n.startswith("xl/worksheets/sheet") for n in changed), changed
    assert set(after) == set(before)
    shown = _displayed_header(path)
    assert any(isinstance(v, str) and CLIENT in v for v in shown.values()), f"{code} 表头没填上编制单位"


@pytest.mark.parametrize(("code", "sheets"), [
    ("A3-8", ("A3-7内部往来核对表", "A3-8商誉减值测试", "A3-8-1可收回金额测试")),
    ("E0", ()),
])
async def test_real_non_active_sheet_headers_show_project_values(tmp_path, monkeypatch, code, sheets):
    """活动页是「编制说明」（A3-7 / A3-8）或各页自带表头（E0）：其余各页表头不再显示示例，点名的页能看到编制单位
    （A3-8-1 的表头是引用 A3-8 的公式，按缓存值读也要显示）。"""
    path = _generate(tmp_path, monkeypatch, code)
    res = await _fill(path, code=code)
    assert res["status"] == "ok", res
    shown = _displayed_header(path)
    labels = {k: v for k, v in shown.items() if isinstance(v, str) and hdr._header_label(v)}
    traces = {k: v for k, v in labels.items() if _EXAMPLE_TRACE.search(v)}
    assert not traces, f"{code} 表头标签仍显示模板示例：{list(traces.items())[:5]}"
    for title in sheets:
        assert any(s == title and isinstance(v, str) and CLIENT in v for (s, _c), v in shown.items()), (
            f"{code}「{title}」表头看不到编制单位")


@pytest.mark.parametrize("code", ["A3-8", "E0"])
async def test_real_external_data_refs_in_header_area_are_untouched(tmp_path, monkeypatch, code):
    """表头区里缓存是数据的外部引用（A3-8 可收回金额、E0 银行名称）不是表头：公式与缓存值都不动。"""
    path = _generate(tmp_path, monkeypatch, code)
    data_before = _external_data_refs(path)
    assert data_before, f"{code} 应有缓存为数据的外部引用（模板变了？）"
    await _fill(path, code=code)
    assert _external_data_refs(path) == data_before


@pytest.mark.parametrize("code", ["D2-5", "D4-12", "D4-21", "D4-5", "E1-28", "J2A", "A1-14"])
async def test_real_external_link_header_shows_project_values(tmp_path, monkeypatch, code):
    """表头引用了其他工作簿的模板（spec post-push-residual-defect-closure R3）：生成后各页表头不再显示示例、
    活动页能看到编制单位（按缓存值读，与 HTML 渲染同口径）；缓存是数据的外部引用原样不动。
    原先这类模板回报「外部链接，请手工填写」，界面上是「会计期间：202X年度」「编制人：张XX」。"""
    path = _generate(tmp_path, monkeypatch, code)
    data_before = _external_data_refs(path)
    assert any(hdr._EXTERNAL_SINGLE_REF_RE.match(str(v)) for v in _snapshot(openpyxl.load_workbook(path)).values()), (
        f"{code} 表头应有外部链接引用（模板变了？）")
    res = await _fill(path, code=code)
    assert res["status"] == "ok", res
    assert any(f["field"].startswith("external_") for f in res["filled"]), f"{code} 没有改写外部引用"
    shown = _displayed_header(path)
    traces = {k: v for k, v in shown.items() if isinstance(v, str) and _EXAMPLE_TRACE.search(v)}
    assert not traces, f"{code} 表头仍显示模板示例：{list(traces.items())[:5]}"
    wb = openpyxl.load_workbook(path)
    active = wb.active.title
    wb.close()
    assert any(s == active and isinstance(v, str) and CLIENT in v for (s, _c), v in shown.items()), (
        f"{code} 活动页「{active}」表头看不到编制单位")
    assert _external_data_refs(path) == data_before, f"{code} 缓存是数据的外部引用被改了"

"""Wave 2 守卫：自定义底稿单元格编辑端点 + 存量投影补齐。

spec: .kiro/specs/custom-workpaper-dual-mode-formula-and-batch/
  Task 6（单元格编辑端点）/ Task 27（存量 render 补齐）/ Task 8（本守卫）

覆盖 Property：
  1  xlsx 权威不变式（先写 xlsx 再刷投影，不存在只写投影的写入路径）
  12 零回归（custom 门控存在 ⇒ 其余 componentType 路径不受影响）
  13 / 22 存量底稿 render 时补齐（幂等 + 不污染非 custom）

🔴 变异检验（比对失败测试名集合，不看退出码）：
  ① 端点里把 write_cells_to_xlsx / refresh_custom_projection 两步顺序调换
  ② 非 custom 的 409 改成静默放行
  ③ Task 27 的补齐去掉 `component_type == "custom"` 门控
  ④ project_if_empty 改成无条件重投影（破坏幂等短路）
  ⑤ 去掉 MAX_CELL_UPDATES 超限分支
"""

from __future__ import annotations

import re
from pathlib import Path
from typing import Any

import pytest

from app.services import custom_workpaper_projection as proj
from app.services.custom_workpaper_projection import (
    grid_has_content,
    project_custom_workpaper,
    project_if_empty,
)

# ─── helpers（与 test_custom_workpaper_projection.py 同款，守卫基建允许各自持有）───

_BACKEND = Path(__file__).resolve().parents[1]
_REPO = _BACKEND.parent


def _strip_comments(src: str) -> str:
    """剥 Python 注释与字符串字面量（含三引号 docstring）。

    🔴 必需：本 spec 的生产代码注释里大量引用被禁/被要求的符号名（解释顺序与门控理由），
    不剥会把说明文字数成真实调用 —— 那会让「调用被删」这个核心变异静默逃逸。
    """
    src = re.sub(r'"""[\s\S]*?"""', '""', src)
    src = re.sub(r"'''[\s\S]*?'''", "''", src)
    src = re.sub(r"(?m)#.*$", "", src)
    return src


def _read(rel: str) -> str:
    p = _REPO / rel
    assert p.exists(), f"路径不存在（守卫判据失效）: {rel}"
    txt = p.read_text(encoding="utf-8")
    assert txt.strip(), f"文件为空（守卫判据失效）: {rel}"
    return txt


def _func_body(src: str, name: str) -> str:
    """截取 `def name(` / `async def name(` 的函数体。

    🔴 缩进正则用 `[ \\t]*` 不用 `\\s*`（后者在 re.M 下含换行，会从空行开始匹配、
       把函数体截成空串）；🔴 先用圆括号配对跳过多行参数列表，否则收尾 `):` 那行
       缩进为 0 会让扫描立刻 break（平台已记的两族坑）。
    """
    m = re.search(rf"(?m)^([ \t]*)(?:async\s+)?def\s+{re.escape(name)}\s*\(", src)
    assert m, f"未找到函数 {name}（守卫判据失效）"
    indent = len(m.group(1))
    i = src.index("(", m.end() - 1)
    depth = 0
    while i < len(src):
        if src[i] == "(":
            depth += 1
        elif src[i] == ")":
            depth -= 1
            if depth == 0:
                break
        i += 1
    assert depth == 0, f"{name} 参数列表括号不配对（守卫判据失效）"
    lines = src[i:].splitlines()
    body = [src[m.start():i].replace("\n", " ")]
    for ln in lines[1:]:
        if ln.strip() and (len(ln) - len(ln.lstrip())) <= indent:
            break
        body.append(ln)
    out = "\n".join(body)
    assert len(out) > 50, f"{name} 函数体过短（截取失败）"
    return out


_CELLS_ROUTER = "backend/app/routers/custom_workpaper_cells.py"
_RENDER_CONFIG = "backend/app/routers/wp_render_config.py"
_CONTEXT_MOD = "backend/app/services/custom_workpaper_context.py"
_GRID_SHEET = "audit-platform/frontend/src/components/workpaper/GtGridSheet.vue"


@pytest.fixture(autouse=True)
def _no_flag_modified(monkeypatch):
    monkeypatch.setattr(proj, "flag_modified", lambda *_a, **_k: None)


class _FakeWp:
    def __init__(self, file_path: str | None = None, parsed_data: dict | None = None):
        self.file_path = file_path
        self.parsed_data = parsed_data


def _make_xlsx(tmp_path: Path, sheet: str, cells: dict[str, Any]) -> Path:
    import openpyxl

    wb = openpyxl.Workbook()
    ws = wb.active
    ws.title = sheet
    for ref, val in cells.items():
        ws[ref] = val
    fp = tmp_path / "wp.xlsx"
    wb.save(str(fp))
    wb.close()
    return fp


# ─── Property 1: xlsx 权威 —— 端点必须先写 xlsx 再刷投影 ─────────────────────


class TestEndpointWriteOrder:
    def test_both_calls_present(self):
        """反向自检：两个调用都在，否则下面的顺序断言是空转。"""
        body = _strip_comments(_func_body(_read(_CELLS_ROUTER), "update_custom_cells"))
        assert "write_cells_to_xlsx(" in body, "端点未调 write_cells_to_xlsx"
        assert "refresh_custom_projection(" in body, "端点未调 refresh_custom_projection"

    def test_xlsx_write_precedes_projection_refresh(self):
        """🔴 顺序反了会让投影领先于权威：xlsx 写失败时用户看到「界面有、文件没有」的值。"""
        body = _strip_comments(_func_body(_read(_CELLS_ROUTER), "update_custom_cells"))
        i_xlsx = body.index("write_cells_to_xlsx(")
        i_proj = body.index("refresh_custom_projection(")
        assert i_xlsx < i_proj, (
            "写 xlsx 必须在刷投影之前（xlsx 是唯一权威）；"
            f"实测 write_cells_to_xlsx@{i_xlsx} > refresh_custom_projection@{i_proj}"
        )

    def test_file_version_incremented(self):
        """与既有 univer-save 的版本语义一致（R3.3）。"""
        body = _strip_comments(_func_body(_read(_CELLS_ROUTER), "update_custom_cells"))
        assert re.search(r"file_version\s*=", body), "写入后未递增 file_version"

    def test_xlsx_write_failure_not_swallowed(self):
        """写盘失败必须让请求失败 —— xlsx 是权威，写不进去不能报成功。"""
        body = _strip_comments(_func_body(_read(_CELLS_ROUTER), "update_custom_cells"))
        assert "HTTPException" in body, "端点无 HTTPException（错误无法上报）"
        assert not re.search(r"except\s+Exception[^\n]*:\s*\n\s*pass", body), (
            "写 xlsx 的异常被 pass 吞掉 ⇒ 会向用户报成功"
        )


class TestUpdateLimits:
    def test_max_cell_updates_constant(self):
        from app.routers.custom_workpaper_cells import MAX_CELL_UPDATES

        assert MAX_CELL_UPDATES == 500

    def test_overflow_branch_exists(self):
        """超限必须 422 + `overflow` 标记，不得静默截断（截断让用户以为都存上了）。

        🔴 判据必须断言**条件形态**而非字符串存在 —— 变异实测：把
        `if len(updates) > MAX_CELL_UPDATES:` 改成 `if False:` 后，
        「`MAX_CELL_UPDATES` 在函数体里」与「`overflow` 在函数体里」两条断言**仍然通过**
        （常量还在类型注解/注释外的其他位置、`overflow` 还在字典里），
        即最核心的变异静默逃逸（平台已记的「弱判据」同族坑）。
        """
        body = _strip_comments(_func_body(_read(_CELLS_ROUTER), "update_custom_cells"))
        assert re.search(
            r"if\s+len\(\s*updates\s*\)\s*>\s*MAX_CELL_UPDATES\s*:", body
        ), "端点未按 `len(updates) > MAX_CELL_UPDATES` 校验单次写入上限"
        assert "overflow" in body, "超限响应缺 overflow 标记"
        assert re.search(r"status_code\s*=\s*422", body), "超限未返回 422"

    def test_illegal_ref_rejected_before_touching_xlsx(self):
        """引用非法要在碰 xlsx 之前拒绝，否则会写进去一半再报错。"""
        body = _strip_comments(_func_body(_read(_CELLS_ROUTER), "update_custom_cells"))
        i_norm = body.index("normalize_cell_ref(")
        i_xlsx = body.index("write_cells_to_xlsx(")
        assert i_norm < i_xlsx, "单元格引用校验必须早于 xlsx 写入"


# ─── 非 custom 底稿必须被拒（409），不得静默写坏标准底稿 ─────────────────────


class TestCustomGating:
    def test_both_endpoints_go_through_gate(self):
        src = _strip_comments(_read(_CELLS_ROUTER))
        for fn in ("update_custom_cells", "refresh_custom_projection_endpoint"):
            body = _func_body(src, fn)
            assert "_require_custom_ctx(" in body, f"{fn} 未经 custom 校验闸"

    def test_gate_returns_409_for_non_custom(self):
        body = _strip_comments(_func_body(_read(_CELLS_ROUTER), "_require_custom_ctx"))
        assert "is_custom" in body, "闸未检查 is_custom"
        assert "409" in body, "非 custom 底稿未返回 409（会静默写坏标准底稿的 xlsx）"

    def test_gate_returns_404_and_403(self):
        body = _strip_comments(_func_body(_read(_CELLS_ROUTER), "_require_custom_ctx"))
        assert "404" in body, "底稿不存在未返回 404"
        assert "403" in body, "只读态未返回 403"

    def test_refresh_projection_does_not_require_writable(self):
        """重投影不改 xlsx，归档底稿也应能刷出正确内容。"""
        body = _strip_comments(
            _func_body(_read(_CELLS_ROUTER), "refresh_custom_projection_endpoint")
        )
        assert "require_writable=False" in body, (
            "重投影端点不应要求可写（它不改 xlsx）"
        )


class TestCustomContextResolver:
    """`resolve_is_custom` 是「这张底稿是不是 custom」的单一真源。"""

    def test_read_only_statuses_only_archived(self):
        from app.models.workpaper_models import WpFileStatus
        from app.services.custom_workpaper_context import READ_ONLY_FILE_STATUSES

        assert READ_ONLY_FILE_STATUSES == frozenset({WpFileStatus.archived})

    def test_uses_same_predicates_as_render_config(self):
        """与 `_maybe_custom_classifications` 同一套谓词（否则判定漂移）。"""
        ctx_src = _strip_comments(_read(_CONTEXT_MOD))
        assert "is_standard_wp_code" in ctx_src
        assert "ProcedureInstance" in ctx_src
        assert "WpSourceType.manual" in ctx_src

    def test_does_not_import_concurrently_edited_helpers(self):
        """🔴 有意不 import `wp_render_config_helpers`（它正被并发 spec 高频改动）。"""
        ctx_src = _strip_comments(_read(_CONTEXT_MOD))
        assert "wp_render_config_helpers" not in ctx_src, (
            "不得依赖 wp_render_config_helpers（并发改动热点），"
            "应直接复用 acnr.grammar.is_standard_wp_code"
        )

    @pytest.mark.parametrize(
        ("wp_code", "expected"),
        [("MYCUSTOM1", True), ("D2-1", False), ("K1", False), (None, False), ("", False)],
    )
    def test_manual_non_standard_code_is_custom(self, wp_code, expected):
        """手工建 + 编号不像标准编号 ⇒ custom；标准编号 ⇒ 不是（拒绝写入）。"""
        import asyncio

        from app.models.workpaper_models import WpSourceType
        from app.services.custom_workpaper_context import resolve_is_custom

        class _Db:
            async def execute(self, *_a, **_k):
                class _R:
                    @staticmethod
                    def scalar():
                        return 0  # 无自定义程序实例

                return _R()

        class _Wp:
            id = "wp"
            project_id = "p"
            source_type = WpSourceType.manual

        got = asyncio.run(resolve_is_custom(_Db(), _Wp(), wp_code))
        assert got is expected

    def test_query_failure_refuses_rather_than_allows(self):
        """判不出来返回 False（拒绝写入），而不是 True（放行写坏文件）。"""
        import asyncio

        from app.models.workpaper_models import WpSourceType
        from app.services.custom_workpaper_context import resolve_is_custom

        class _Db:
            async def execute(self, *_a, **_k):
                raise RuntimeError("boom")

        class _Wp:
            id = "wp"
            project_id = "p"
            source_type = WpSourceType.manual

        # 编号不像标准编号，但查询已炸 → 必须 False（fail-closed 方向）
        assert asyncio.run(resolve_is_custom(_Db(), _Wp(), "MYCUSTOM1")) is False


# ─── Property 13 / 22: 存量底稿 render 时补齐 ────────────────────────────────


class TestRenderBackfillWiring:
    def test_render_config_calls_project_if_empty(self):
        src = _strip_comments(_read(_RENDER_CONFIG))
        assert "project_if_empty(" in src, "render-config 未接存量投影补齐（R2.6）"

    def test_backfill_is_gated_on_custom(self):
        """🔴 无门控会污染其余 37 种 componentType 的 render 输出（Property 12）。"""
        src = _strip_comments(_read(_RENDER_CONFIG))
        idx = src.index("project_if_empty(")
        window = src[max(0, idx - 600):idx]
        assert re.search(r'component_type\s*==\s*"custom"', window), (
            "存量补齐必须在 `component_type == \"custom\"` 门控内，"
            "否则会对全部 componentType 生效"
        )

    def test_backfill_is_fail_open(self):
        """补齐失败不得阻断整个 render-config。"""
        src = _strip_comments(_read(_RENDER_CONFIG))
        idx = src.index("project_if_empty(")
        window = src[max(0, idx - 400):idx + 600]
        assert "try:" in window and "except" in window, "补齐未 fail-open"

    def test_render_config_does_not_persist_on_get(self):
        """🔴 render-config 是 GET，补齐不落库（避免读端点写库 / 锁竞争）。"""
        src = _strip_comments(_read(_RENDER_CONFIG))
        idx = src.index("project_if_empty(")
        window = src[idx:idx + 500]
        assert "write_projection_to_parsed_data" not in window, (
            "GET render-config 内不应写库（补齐是纯计算，用户编辑时才持久化）"
        )


class TestProjectIfEmptyBehavior:
    def test_short_circuits_when_content_present(self):
        """幂等短路：已有内容时原样返回（不重读 xlsx）。"""
        existing = {"cells": {"B6": {"v": 1}}, "max_row": 6, "max_col": 2}
        got = project_if_empty(_FakeWp(file_path="/nope/none.xlsx"), "X9", existing)
        assert got is existing

    def test_backfills_from_xlsx_when_empty(self, tmp_path):
        fp = _make_xlsx(tmp_path, "X9", {"A1": "致同", "B6": 123.45})
        got = project_if_empty(_FakeWp(file_path=str(fp)), "X9", None)
        assert grid_has_content(got), "空投影未被补齐"
        assert "B6" in got["cells"], "补齐后坐标非恒等（B6 丢失）"
        assert got["cells"]["B6"]["v"] == 123.45

    def test_idempotent(self, tmp_path):
        fp = _make_xlsx(tmp_path, "X9", {"B6": 1})
        first = project_if_empty(_FakeWp(file_path=str(fp)), "X9", None)
        second = project_if_empty(_FakeWp(file_path=str(fp)), "X9", first)
        assert second is first, "已有内容时应短路（否则每次 render 都读 xlsx）"

    def test_empty_dict_and_zero_maxrow_treated_as_empty(self, tmp_path):
        """`{}` 与 `{cells:{...}, max_row:0}` 都算无内容（与前端 hasData 同口径）。"""
        fp = _make_xlsx(tmp_path, "X9", {"B6": 1})
        wp = _FakeWp(file_path=str(fp))
        for existing in ({}, {"cells": {}}, {"cells": {"B6": {"v": 1}}, "max_row": 0}):
            got = project_if_empty(wp, "X9", existing)
            assert grid_has_content(got), f"未被判为空: {existing}"

    def test_source_unavailable_distinguishes_missing_file(self, tmp_path):
        """🔴 「文件损坏」与「空底稿」必须可区分（R1.4），否则损坏被当正常空表。"""
        missing = project_if_empty(_FakeWp(file_path=str(tmp_path / "no.xlsx")), "X9", None)
        assert missing.get("source_unavailable") is True

        empty_fp = _make_xlsx(tmp_path, "X9", {})
        present = project_if_empty(_FakeWp(file_path=str(empty_fp)), "X9", None)
        assert present.get("source_unavailable") is False

    def test_backfill_keeps_six_key_contract(self, tmp_path):
        required = {"cells", "max_row", "max_col", "col_widths", "merged_cells", "header_rows"}
        for wp in (_FakeWp(file_path=None), _FakeWp(file_path=str(tmp_path / "x.xlsx"))):
            got = project_if_empty(wp, "X9", None)
            assert required <= set(got), f"fail-open 路径键集不全: {sorted(got)}"


class TestGridHasContentParityWithFrontend:
    """跨前后端交叉锁死：后端 `grid_has_content` 必须与前端 `hasData` 同口径。

    🔴 两侧不一致会让补齐白跑 —— 后端判「有内容」不再补，前端判「无内容」显示空态。
    """

    def test_frontend_hasdata_expression_unchanged(self):
        vue = _read(_GRID_SHEET)
        m = re.search(r"const\s+hasData\s*=\s*computed\(\(\)\s*=>\s*([^\n]+)", vue)
        assert m, "未找到 GtGridSheet.hasData 定义（守卫判据失效）"
        expr = m.group(1)
        assert "length > 0" in expr, f"前端 hasData 口径已变: {expr}"
        assert re.search(r"maxRow(?:\.value)?\s*>\s*0", expr), f"前端 hasData 口径已变: {expr}"

    @pytest.mark.parametrize(
        ("grid", "expected"),
        [
            (None, False),
            ({}, False),
            ({"cells": {}, "max_row": 5}, False),
            ({"cells": {"A1": {}}, "max_row": 0}, False),
            ({"cells": {"A1": {}}}, False),
            ({"cells": {"A1": {}}, "max_row": 1}, True),
            ({"cells": {"A1": {}}, "max_row": "3"}, True),
            ({"cells": {"A1": {}}, "max_row": "x"}, False),
        ],
    )
    def test_backend_matches(self, grid, expected):
        assert grid_has_content(grid) is expected


class TestRouterRegistered:
    def test_registered_in_router_registry(self):
        """新 router 必须挂进 registry，否则端点 404 而前端只见「保存失败」。"""
        src = _read("backend/app/router_registry/workpaper.py")
        assert "custom_workpaper_cells" in src, "router 未在 registry 注册"
        stripped = _strip_comments(src)
        assert re.search(
            r"from app\.routers\.custom_workpaper_cells import router as custom_workpaper_cells",
            stripped,
        ), "registry 未 import 该 router"
        assert re.search(r'"渲染":\s*\[[^\]]*custom_workpaper_cells', stripped), (
            "router 未加入分组列表（import 了但没挂 = 端点仍 404）"
        )

    def test_routes_exposed(self):
        from app.routers.custom_workpaper_cells import router

        paths = {r.path for r in router.routes}
        assert "/api/workpapers/{wp_id}/custom-cells" in paths
        assert "/api/workpapers/{wp_id}/custom-refresh-projection" in paths


class TestNoProjectionOnlyWritePath:
    """Property 1 收口：不存在「只写投影不写 xlsx」的**新增**生产路径。

    🔴 既有 `write_cell_to_parsed_data` 是历史单写投影路径，由 Task 13 给它补 xlsx 双写；
    本守卫只钉住本 spec 新增的 cells 端点不得退化成单写。
    """

    def test_cells_endpoint_never_writes_projection_without_xlsx(self):
        body = _strip_comments(_func_body(_read(_CELLS_ROUTER), "update_custom_cells"))
        assert "write_projection_to_parsed_data(" not in body, (
            "端点不应直接写投影 —— 必须经 refresh_custom_projection 从 xlsx 重投，"
            "直写投影等于绕过权威"
        )

    def test_projection_module_documents_authority(self):
        doc = proj.__doc__ or ""
        assert "xlsx" in doc and "权威" in doc, "投影模块 docstring 未写明 xlsx 权威口径"


# ════════════════════════════════════════════════════════════════════════════
# Wave 4 Task 15：公式求值双写 / 删公式清格（Property 8 / Property 21）
#
# 这里做**运行时**验证（真实 xlsx 往返），不只做源码级断言 ——
# 源码级只能证明「调用存在且顺序对」，证明不了「xlsx 与投影双方真的都有值」。
# ════════════════════════════════════════════════════════════════════════════


class TestFormulaDualWriteRuntime:
    """Property 8：公式求值结果必须同时落 xlsx 与投影。"""

    def test_value_lands_in_both_xlsx_and_projection(self, tmp_path):
        from app.services.custom_workpaper_projection import (
            refresh_custom_projection,
            write_cells_to_xlsx,
        )

        sheet = "GT-CUSTOM-1"
        fp = _make_xlsx(tmp_path, sheet, {"A1": "项目", "B1": "金额"})
        wp = _FakeWp(file_path=str(fp), parsed_data={})

        # 模拟 save_formula 的 custom 分支：先写 xlsx（权威），再刷投影
        written = write_cells_to_xlsx(str(fp), sheet, {"B5": "1234.50"})
        assert written == 1
        grid = refresh_custom_projection(wp, sheet)

        # ① 投影侧有值
        assert "B5" in grid["cells"], "投影里应有 B5"
        assert str(grid["cells"]["B5"]["v"]) == "1234.50"

        # ② xlsx 本体也有值（重新独立读一遍，不信任上一步的返回）
        import openpyxl

        wb = openpyxl.load_workbook(str(fp), data_only=False)
        try:
            assert wb[sheet]["B5"].value == "1234.50", "xlsx 本体里应有 B5"
        finally:
            wb.close()

        # ③ parsed_data 里的投影与返回值一致（写回真的落到 wp 上）
        assert wp.parsed_data["html_data"][sheet]["cells"]["B5"]["v"] == "1234.50"

    def test_projection_is_reprojected_not_patched(self, tmp_path):
        """投影必须**从 xlsx 重投影**而非就地打补丁。

        🔴 就地打补丁会让 xlsx 里已被删除的格在投影里残留 ——
        投影必须恒等于权威。
        """
        from app.services.custom_workpaper_projection import (
            refresh_custom_projection,
            write_cells_to_xlsx,
        )

        sheet = "GT-CUSTOM-1"
        fp = _make_xlsx(tmp_path, sheet, {"A1": "项目", "B5": "old"})
        wp = _FakeWp(file_path=str(fp), parsed_data={})
        refresh_custom_projection(wp, sheet)
        assert wp.parsed_data["html_data"][sheet]["cells"]["B5"]["v"] == "old"

        # 直接在 xlsx 里清掉 B5，再刷投影 → 投影里不应残留
        write_cells_to_xlsx(str(fp), sheet, {"B5": None})
        grid = refresh_custom_projection(wp, sheet)
        assert "B5" not in grid["cells"], "重投影后不应残留已删除的格"

    def test_xlsx_write_failure_raises_so_projection_never_leads(self, tmp_path):
        """xlsx 写失败必须抛异常 —— 否则会留下「只有投影有值」的分叉态。"""
        from app.services.custom_workpaper_projection import write_cells_to_xlsx

        missing = tmp_path / "__not_here__.xlsx"
        with pytest.raises(FileNotFoundError):
            write_cells_to_xlsx(str(missing), "GT-CUSTOM-1", {"B5": 1})

        # sheet 不存在同样抛（不静默创建）
        fp = _make_xlsx(tmp_path, "GT-CUSTOM-1", {"A1": "项目"})
        with pytest.raises(KeyError):
            write_cells_to_xlsx(str(fp), "NO-SUCH-SHEET", {"B5": 1})


class TestFormulaDeleteClearsCellRuntime:
    """Property 21：删公式后 xlsx 与投影双方该格都必须为空。"""

    def test_clear_removes_value_from_both_sides(self, tmp_path):
        from app.services.custom_workpaper_projection import (
            refresh_custom_projection,
            write_cells_to_xlsx,
        )

        sheet = "GT-CUSTOM-1"
        fp = _make_xlsx(tmp_path, sheet, {"A1": "项目"})
        wp = _FakeWp(file_path=str(fp), parsed_data={})

        # 先落一个"公式求值结果"
        write_cells_to_xlsx(str(fp), sheet, {"C7": "999.00"})
        refresh_custom_projection(wp, sheet)
        assert wp.parsed_data["html_data"][sheet]["cells"]["C7"]["v"] == "999.00"

        # 删公式 → 清格（写 None）+ 刷投影
        write_cells_to_xlsx(str(fp), sheet, {"C7": None})
        grid = refresh_custom_projection(wp, sheet)

        assert "C7" not in grid["cells"], "投影里该格应已清空"
        import openpyxl

        wb = openpyxl.load_workbook(str(fp), data_only=False)
        try:
            assert wb[sheet]["C7"].value is None, "xlsx 里该格应已清空"
        finally:
            wb.close()

    def test_cleared_cell_is_editable_again(self, tmp_path):
        """清格后该格恢复可手工编辑（写入新值应成功且双侧一致）。"""
        from app.services.custom_workpaper_projection import (
            refresh_custom_projection,
            write_cells_to_xlsx,
        )

        sheet = "GT-CUSTOM-1"
        fp = _make_xlsx(tmp_path, sheet, {"A1": "项目"})
        wp = _FakeWp(file_path=str(fp), parsed_data={})
        write_cells_to_xlsx(str(fp), sheet, {"C7": None})
        write_cells_to_xlsx(str(fp), sheet, {"C7": "手工值"})
        grid = refresh_custom_projection(wp, sheet)
        assert grid["cells"]["C7"]["v"] == "手工值"


class TestFormulaRouterGating:
    """Property 12 零回归：非 custom 底稿的 save/delete 路径不触发 xlsx 写入。"""

    _SRC = None

    @classmethod
    def _src(cls) -> str:
        if cls._SRC is None:
            cls._SRC = _strip_comments(_read("backend/app/routers/wp_formula.py"))
        return cls._SRC

    @staticmethod
    def _call_pos(body: str, name: str) -> int:
        """调用点位置（-1 表示无调用）。

        🔴 必须匹配 `name(` 而非裸 `name` —— 函数体内的**局部 import**
        （`from ... import (refresh_custom_projection, write_cells_to_xlsx,)`）
        也含该标识符，裸 find 会命中 import 行 ⇒ 「调用被删」这个核心变异
        静默逃逸（本守卫首版即如此，变异 B1/B3 实测 GREEN）。
        """
        m = re.search(rf"{re.escape(name)}\s*\(", body)
        return m.start() if m else -1

    def test_save_formula_xlsx_write_is_inside_custom_gate(self):
        body = _func_body(self._src(), "save_formula")
        gate = body.find("if _is_custom_wp:")
        xlsx = self._call_pos(body, "write_cells_to_xlsx")
        assert gate > 0, "save_formula 应有 custom 门控"
        assert xlsx >= 0, "save_formula 必须真正调用 write_cells_to_xlsx（不能只 import）"
        assert xlsx > gate, "xlsx 写入必须在 custom 门控之内（非 custom 不得写 xlsx）"

    def test_save_formula_keeps_projection_only_path_for_non_custom(self):
        body = _func_body(self._src(), "save_formula")
        assert "write_cell_to_parsed_data(" in body, "非 custom 分支应保留原写投影路径"

    def test_delete_formula_xlsx_write_is_inside_custom_gate(self):
        body = _func_body(self._src(), "delete_formula")
        gate = self._call_pos(body, "resolve_is_custom")
        xlsx = self._call_pos(body, "write_cells_to_xlsx")
        assert gate >= 0, "delete_formula 应有 custom 门控"
        assert xlsx >= 0, "delete_formula 必须真正调用 write_cells_to_xlsx 清格"
        assert xlsx > gate, "清格必须在 custom 门控之内"

    def test_delete_formula_clears_with_none(self):
        """清格必须写 None（写空串只会把该格变成空文本，不是清空）。"""
        body = _func_body(self._src(), "delete_formula")
        assert re.search(
            r"write_cells_to_xlsx\([^)]*\{[^}]*:\s*None\s*\}", body, re.S
        ), "清格应写 None"

    def test_delete_clear_failure_does_not_fail_the_delete(self):
        body = _func_body(self._src(), "delete_formula")
        assert "cell_clear_failed" in body, "清格失败应有可见标记"
        idx = body.find("cell_clear_failed")
        around = body[max(0, idx - 500):idx + 200]
        assert "except" in around, "清格失败应被捕获"
        assert "raise HTTPException" not in around, "清格失败不得让删除整体失败"

    def test_resolve_is_custom_called_with_three_args(self):
        """🔴 少传一个参数会 TypeError 被 except 吞成死代码（H 循环 Task 5 同型）。"""
        for m in re.finditer(r"resolve_is_custom\(([^)]*)\)", self._src()):
            args = [a.strip() for a in m.group(1).split(",") if a.strip()]
            assert len(args) == 3, f"resolve_is_custom 应传 3 参，实为 {args}"

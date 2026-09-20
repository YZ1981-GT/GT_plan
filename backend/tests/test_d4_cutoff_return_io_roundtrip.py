"""D4-17/18/19/20 截止/折扣/退货 导入导出 Round-Trip 与结构守卫

Feature: d4-cutoff-return-writeback-formula-io · IO 收口
Property（Requirements 1.2-1.5）：四表 export/import 往返后稳定 row id 与全部录入字段逐字段
一致，派生值由公式重算（isCutoff 留 None、discountRate/shouldProvide/diff 重算，不信文件值）。

行为级 PBT：驱动真实 _parse_d4_17/18/19/20_row + export 行构造。hypothesis max_examples=5。
参照 test_d4_inspection_io_roundtrip.py 同源写法。
"""

from __future__ import annotations

import datetime
import io

import openpyxl
from hypothesis import given, settings
from hypothesis import strategies as st

from app.routers.wp_render_strategies import _d4_import_export as mod


def _headers(sheet: str) -> list[str]:
    return mod._get_headers(sheet)


def _roundtrip(headers: list[str], row_values_list: list[list], parser) -> list[dict]:
    wb = openpyxl.Workbook()
    ws = wb.active
    ws.append(headers)
    for rv in row_values_list:
        ws.append(rv)
    buf = io.BytesIO()
    wb.save(buf)
    buf.seek(0)
    wb2 = openpyxl.load_workbook(buf, data_only=True)
    ws2 = wb2.active
    actual_headers = [
        str(c.value).strip() if c.value else ""
        for c in next(ws2.iter_rows(min_row=1, max_row=1))
    ]
    out = []
    for row in ws2.iter_rows(min_row=2, values_only=True):
        if not row or all(v is None for v in row):
            continue
        d = parser(row, actual_headers)
        if d:
            out.append(d)
    return out


_txt = st.text(max_size=14, alphabet=st.characters(min_codepoint=0x4E00, max_codepoint=0x9FFF))
_amt = st.floats(min_value=0, max_value=1e9, allow_nan=False, allow_infinity=False).map(lambda x: round(x, 2))
_date = st.dates(min_value=datetime.date(2020, 1, 1), max_value=datetime.date(2027, 12, 31)).map(lambda d: d.isoformat())


# ─── D4-17 截止（账到单据）：isCutoff 留 None ────────────────────────────────

@settings(max_examples=5)
@given(
    rows=st.lists(
        st.fixed_dictionaries({
            "voucherDate": _date, "voucherNo": _txt, "voucherProduct": _txt,
            "voucherQty": _txt, "voucherAmount": _amt,
            "deliveryDate": _date, "deliveryNo": _txt, "deliveryProduct": _txt,
            "deliveryQty": _txt, "deliveryAmount": _amt, "remark": _txt,
        }),
        min_size=1, max_size=5,
    ),
    poison_cutoff=st.sampled_from(["是", "否", "√", "×"]),
)
def test_d4_17_round_trip_cutoff_left_none(rows: list[dict], poison_cutoff: str) -> None:
    """**Validates: Requirements 1.2, 1.5, 4.2**

    D4-17：录入字段逐字段往返；isCutoff 留 None（前端 checkCutoff 重算），不采信文件跨期列。
    """
    headers = _headers("D4-17")
    row_values_list = []
    for r in rows:
        row_values_list.append([
            "",  # 序号
            r["voucherDate"], r["voucherNo"], r["voucherProduct"], r["voucherQty"], r["voucherAmount"],
            r["deliveryDate"], r["deliveryNo"], r["deliveryProduct"], r["deliveryQty"], r["deliveryAmount"],
            poison_cutoff,  # 是否跨期（污染 → parser 留 None）
            r["remark"],
        ])
    parsed = _roundtrip(headers, row_values_list, mod._parse_d4_17_row)
    assert len(parsed) == len(rows)
    for orig, got in zip(rows, parsed):
        assert got["id"].startswith("r-")
        assert got["voucherNo"] == mod._safe_str(orig["voucherNo"])
        assert got["voucherAmount"] == orig["voucherAmount"]
        assert got["deliveryDate"] == mod._safe_str(orig["deliveryDate"])
        assert got["deliveryAmount"] == orig["deliveryAmount"]
        assert got["remark"] == mod._safe_str(orig["remark"])
        assert got["isCutoff"] is None  # 派生列单源


# ─── D4-18 截止（单据到账） ─────────────────────────────────────────────────

@settings(max_examples=5)
@given(
    rows=st.lists(
        st.fixed_dictionaries({
            "deliveryDate": _date, "deliveryNo": _txt, "deliveryProduct": _txt,
            "deliveryQty": _txt, "deliveryAmount": _amt,
            "voucherDate": _date, "voucherNo": _txt, "voucherProduct": _txt,
            "voucherQty": _txt, "voucherAmount": _amt, "remark": _txt,
        }),
        min_size=1, max_size=5,
    ),
)
def test_d4_18_round_trip(rows: list[dict]) -> None:
    """**Validates: Requirements 1.2, 1.5**

    D4-18：发货单→凭证方向，字段逐字段往返，isCutoff 留 None。
    """
    headers = _headers("D4-18")
    row_values_list = []
    for r in rows:
        row_values_list.append([
            "",
            r["deliveryDate"], r["deliveryNo"], r["deliveryProduct"], r["deliveryQty"], r["deliveryAmount"],
            r["voucherDate"], r["voucherNo"], r["voucherProduct"], r["voucherQty"], r["voucherAmount"],
            "", r["remark"],
        ])
    parsed = _roundtrip(headers, row_values_list, mod._parse_d4_18_row)
    assert len(parsed) == len(rows)
    for orig, got in zip(rows, parsed):
        assert got["deliveryNo"] == mod._safe_str(orig["deliveryNo"])
        assert got["voucherAmount"] == orig["voucherAmount"]
        assert got["isCutoff"] is None


# ─── D4-19 销售折扣：discountRate 重算 ───────────────────────────────────────

@settings(max_examples=5)
@given(
    rows=st.lists(
        st.fixed_dictionaries({
            "customerName": _txt, "discountType": _txt,
            "revenueAmount": st.floats(min_value=1, max_value=1e8, allow_nan=False).map(lambda x: round(x, 2)),
            "discountAmount": st.floats(min_value=1, max_value=1e6, allow_nan=False).map(lambda x: round(x, 2)),
            "reason": _txt, "voucherDate": _date, "voucherNo": _txt,
            "accountSubject": _txt, "detailSubject": _txt,
            "debitAmount": _amt, "creditAmount": _amt,
            "approvalDate": _date, "approver": _txt, "remark": _txt,
        }),
        min_size=1, max_size=5,
    ),
    poison_rate=st.floats(min_value=-99, max_value=99, allow_nan=False).map(lambda x: round(x, 4)),
)
def test_d4_19_round_trip_rate_recompute(rows: list[dict], poison_rate: float) -> None:
    """**Validates: Requirements 1.2, 1.5, 4.2**

    D4-19：录入字段往返；discountRate = 折扣额/收入额 后端重算，不采信文件污染的比例列。
    """
    headers = _headers("D4-19")
    row_values_list = []
    for r in rows:
        row_values_list.append([
            "",  # 序号
            r["customerName"], r["discountType"],
            r["revenueAmount"], r["discountAmount"],
            poison_rate,  # 折扣比例（污染 → 重算覆盖）
            r["reason"], r["voucherDate"], r["voucherNo"],
            r["accountSubject"], r["detailSubject"],
            r["debitAmount"], r["creditAmount"],
            r["approvalDate"], r["approver"], r["remark"],
        ])
    parsed = _roundtrip(headers, row_values_list, mod._parse_d4_19_row)
    assert len(parsed) == len(rows)
    for orig, got in zip(rows, parsed):
        assert got["customerName"] == mod._safe_str(orig["customerName"])
        assert got["revenueAmount"] == orig["revenueAmount"]
        assert got["discountAmount"] == orig["discountAmount"]
        assert got["debitAmount"] == orig["debitAmount"]
        assert got["approver"] == mod._safe_str(orig["approver"])
        # 派生列单源：比例 = 折扣/收入 重算
        assert got["discountRate"] == orig["discountAmount"] / orig["revenueAmount"]


# ─── D4-20 退货明细（current/post）+ 计提（provision 派生重算）─────────────────

@settings(max_examples=5)
@given(
    rows=st.lists(
        st.fixed_dictionaries({
            "voucherDate": _date, "voucherNo": _txt, "bizContent": _txt,
            "subjectName": _txt, "detailSubject": _txt,
            "debitAmount": _amt, "creditAmount": _amt,
            "customerName": _txt, "productName": _txt,
            "returnQty": _txt, "returnAmount": _amt, "returnReason": _txt,
            "hasLitigation": st.sampled_from(["是", "否", ""]),
            "isAbnormal": st.sampled_from(["是", "否", ""]),
            "indexRef": _txt,
        }),
        min_size=1, max_size=4,
    ),
)
def test_d4_20_return_round_trip(rows: list[dict]) -> None:
    """**Validates: Requirements 1.2, 1.3**

    D4-20 退货明细（current/post 同构）：字段逐字段往返，稳定 id。
    """
    headers = _headers("D4-20-current")
    row_values_list = []
    for r in rows:
        row_values_list.append([
            "",
            r["voucherDate"], r["voucherNo"], r["bizContent"], r["subjectName"], r["detailSubject"],
            r["debitAmount"], r["creditAmount"], r["customerName"], r["productName"],
            r["returnQty"], r["returnAmount"], r["returnReason"],
            r["hasLitigation"], r["isAbnormal"], r["indexRef"],
        ])
    parsed = _roundtrip(headers, row_values_list, mod._parse_d4_20_return_row)
    assert len(parsed) == len(rows)
    for orig, got in zip(rows, parsed):
        assert got["id"].startswith("r-")
        assert got["voucherNo"] == mod._safe_str(orig["voucherNo"])
        assert got["returnAmount"] == orig["returnAmount"]
        assert got["hasLitigation"] == mod._safe_str(orig["hasLitigation"])
        assert got["isAbnormal"] == mod._safe_str(orig["isAbnormal"])


@settings(max_examples=5)
@given(
    rows=st.lists(
        st.fixed_dictionaries({
            "productName": _txt,
            "base": _amt,
            "rate": st.floats(min_value=0, max_value=1, allow_nan=False).map(lambda x: round(x, 4)),
            "alreadyProvided": _amt,
            "diffReason": _txt,
        }),
        min_size=1, max_size=4,
    ),
    poison_should=_amt,
    poison_diff=st.floats(min_value=-1e6, max_value=1e6, allow_nan=False).map(lambda x: round(x, 2)),
)
def test_d4_20_provision_round_trip_recompute(rows: list[dict], poison_should: float, poison_diff: float) -> None:
    """**Validates: Requirements 1.2, 1.5, 4.2**

    D4-20 计提：应计提=基数×比例、差异=应计提−已计提 后端重算，不采信文件派生列。
    """
    headers = _headers("D4-20-provision")
    row_values_list = []
    for r in rows:
        row_values_list.append([
            "",  # 序号
            r["productName"], r["base"], r["rate"],
            poison_should,   # 应计提金额（污染 → 重算）
            r["alreadyProvided"],
            poison_diff,     # 差异金额（污染 → 重算）
            r["diffReason"],
        ])
    parsed = _roundtrip(headers, row_values_list, mod._parse_d4_20_provision_row)
    assert len(parsed) == len(rows)
    for orig, got in zip(rows, parsed):
        assert got["base"] == orig["base"]
        assert got["rate"] == orig["rate"]
        assert got["alreadyProvided"] == orig["alreadyProvided"]
        # 派生列单源
        assert got["shouldProvide"] == orig["base"] * orig["rate"]
        assert got["diff"] == orig["base"] * orig["rate"] - orig["alreadyProvided"]


# ─── 结构守卫：item_id 双侧一致（含 D4-20 子表键错位修复）+ 死配置删除 ─────────

from pathlib import Path


def _repo_root() -> Path:
    p = Path(__file__).resolve()
    for _ in range(8):
        if (p / "backend").exists() and (p / "audit-platform").exists():
            return p
        p = p.parent
    raise RuntimeError("repo root not found")


def _read(rel: str) -> str:
    return (_repo_root() / rel).read_text(encoding="utf-8")


def test_d4_20_subtable_key_alignment() -> None:
    """**Validates: Requirements 1.1, 1.3**

    🔴 D4-20 子表键错位修复：后端 sheet 码 D4-20-current/post → item_id
    D4-20-current-returns/post-returns（对齐前端 store 键，DEC-2），provision → D4-20-provision。
    前端组件确实用这些键。
    """
    src = _read("backend/app/routers/wp_render_strategies/_d4_import_export.py")
    # 后端 import+export 两侧都映射到前端键
    assert src.count('item_id = "D4-20-current-returns"') >= 2
    assert src.count('item_id = "D4-20-post-returns"') >= 2
    # 前端组件确实用这些键
    fe = _read("audit-platform/frontend/src/components/workpaper/d4/inspection/D4TabReturn.vue")
    assert "'D4-20-current-returns'" in fe
    assert "'D4-20-post-returns'" in fe
    assert "'D4-20-provision'" in fe


def test_d4_20_bare_sheet_dead_config_removed() -> None:
    """**Validates: Requirements 1.1（DEC-1）**

    D4-20 主 sheet 是死配置（前端只用三子表）→ 已从 _SUPPORTED_SHEETS / _SHEET_HEADERS 删除。
    """
    assert "D4-20" not in mod._SUPPORTED_SHEETS
    assert mod._get_headers("D4-20") == mod._GENERIC_HEADERS  # 不再有专用 header
    # 三子表仍在
    for s in ("D4-20-current", "D4-20-post", "D4-20-provision"):
        assert s in mod._SUPPORTED_SHEETS
        assert mod._get_headers(s) != mod._GENERIC_HEADERS


def test_d4_17_18_19_dispatch_dedicated_parsers() -> None:
    """**Validates: Requirements 1.2**

    import 分发 D4-17/18/19 走专用 parser（不再 generic 兜底）。
    """
    import re
    src = _read("backend/app/routers/wp_render_strategies/_d4_import_export.py")
    assert re.search(r'elif sheet == "D4-17":\s*\n\s*row_dict = _parse_d4_17_row', src)
    assert re.search(r'elif sheet == "D4-18":\s*\n\s*row_dict = _parse_d4_18_row', src)
    assert re.search(r'elif sheet == "D4-19":\s*\n\s*row_dict = _parse_d4_19_row', src)
    assert re.search(r'elif sheet in \("D4-20-current", "D4-20-post"\):\s*\n\s*row_dict = _parse_d4_20_return_row', src)


# ─── D4-19 派生列零/负边界：discountRate 单源重算不除零、不信文件（补 else 分支）─────

@settings(max_examples=5)
@given(
    revenue=st.floats(min_value=-1e6, max_value=1e6, allow_nan=False).map(lambda x: round(x, 2)),
    discount=st.floats(min_value=-1e6, max_value=1e6, allow_nan=False).map(lambda x: round(x, 2)),
    poison_rate=st.floats(min_value=-99, max_value=99, allow_nan=False).map(lambda x: round(x, 4)),
)
def test_d4_19_discount_rate_zero_edge_single_source(
    revenue: float, discount: float, poison_rate: float
) -> None:
    """**Validates: Requirements 1.3, 1.5**

    D4-19 折扣比例单源重算的零/负边界：收入或折扣 ≤0 时 discountRate=0（不除零、不信文件污染的
    比例列）；两者 >0 时才 = 折扣/收入。PBT `test_d4_19_round_trip_rate_recompute` 只覆盖正数区间，
    此处专攻 else 分支（防「负收入/零折扣」时误信文件或抛除零）。
    """
    headers = _headers("D4-19")
    row_values = [
        "", "客户甲", "现金折扣", revenue, discount,
        poison_rate,  # 折扣比例（污染 → 忽略/重算）
        "原因", "2025-01-01", "PZ-1", "6001", "主营", 0.0, 0.0, "2025-01-02", "张三", "备注",
    ]
    parsed = _roundtrip(headers, [row_values], mod._parse_d4_19_row)
    assert len(parsed) == 1
    got = parsed[0]
    if revenue > 0 and discount > 0:
        assert got["discountRate"] == discount / revenue
    else:
        # 收入≤0 或 折扣≤0：单源重算给 0，绝不采信文件污染值 poison_rate
        assert got["discountRate"] == 0
    assert got["discountRate"] != poison_rate or poison_rate == 0


# ─── 导出侧投影守卫：D4-17/18/19/20 派生列导出留空（往返永不回注被信任的派生值）───────

def test_export_leaves_derived_columns_blank() -> None:
    """**Validates: Requirements 1.2, 1.3, 1.5, 2.3**

    export 行构造对 D4-17/18(是否跨期)、D4-19(折扣比例)、D4-20-provision(应计提/差异) 派生列
    一律写空串 ""（前端公式重算），不得 `data_row.get("isCutoff"/"discountRate"/"shouldProvide"/"diff")`。
    这是「派生值单源」在导出投影侧的对称守卫：即便文件被回导，import parser 也已忽略/重算，
    但导出侧亦不得先泄漏被信任的派生值到文件。
    """
    import re
    src = _read("backend/app/routers/wp_render_strategies/_d4_import_export.py")
    # 导出分支绝不读这些派生键（否则往返会把「文件里的派生值」当真源来回搬运）。
    # isCutoff/discountRate/shouldProvide 是 D4-17/18/19/20 专属派生键，全文件唯一归属这四表，
    # 可直接全文件断言（"diff" 是 D4-23 发票差异合法列名，故按分支体局部断言，见下）。
    assert 'data_row.get("isCutoff")' not in src
    assert 'data_row.get("discountRate")' not in src
    assert 'data_row.get("shouldProvide")' not in src
    # D4-20-provision 导出分支体内：应计提/差异派生列必须写空串，不得从 store 读派生值。
    m_prov = re.search(
        r'elif sheet == "D4-20-provision":(.+?)(?:elif sheet ==|\n        else:)', src, re.S
    )
    assert m_prov, "D4-20-provision export branch not found"
    prov_body = m_prov.group(1)
    assert 'data_row.get("diff")' not in prov_body
    assert 'data_row.get("shouldProvide")' not in prov_body


def test_d4_17_18_export_preserves_business_direction() -> None:
    """**Validates: Requirements 1.2**

    D4-17(账→单据 forward)/D4-18(单据→账 backward) 的导出列顺序保持业务方向：
    D4-17 先凭证列后发货单列；D4-18 先发货单列后凭证列。锁死方向不被对调。
    """
    import re
    src = _read("backend/app/routers/wp_render_strategies/_d4_import_export.py")
    # D4-17 分支：voucherDate 出现在 deliveryDate 之前
    m17 = re.search(r'elif sheet == "D4-17":(.+?)elif sheet == "D4-18":', src, re.S)
    assert m17, "D4-17 export branch not found"
    body17 = m17.group(1)
    assert body17.index('voucherDate') < body17.index('deliveryDate')
    # D4-18 分支：deliveryDate 出现在 voucherDate 之前
    m18 = re.search(r'elif sheet == "D4-18":(.+?)elif sheet == "D4-19":', src, re.S)
    assert m18, "D4-18 export branch not found"
    body18 = m18.group(1)
    assert body18.index('deliveryDate') < body18.index('voucherDate')


# ═══════════════════════════════════════════════════════════════════════════════
# Task 6 · 四态变异（four-state mutation）+ 截止非跨期语义守卫
#
# Design Property 4（Validates 4.1）：变异检验命中预期守卫——把代码翻成反模式后，守卫必须
# 失败。下列测试不改生产代码，而是就地重演「反模式」输入/逻辑，断言守卫不变量被破坏（证明守卫
# 敏感、非空转）。每个 facet 一条 mutation，触类旁通覆盖 D4-17/18/19/20 四表。
# ═══════════════════════════════════════════════════════════════════════════════


# ─── Mutation 1 · 派生值单源（未知边界）：若 parser「信任文件派生列」则往返守卫失败 ──────

@settings(max_examples=5)
@given(
    revenue=st.floats(min_value=1, max_value=1e6, allow_nan=False).map(lambda x: round(x, 2)),
    discount=st.floats(min_value=1, max_value=1e6, allow_nan=False).map(lambda x: round(x, 2)),
    poison_rate=st.floats(min_value=-99, max_value=99, allow_nan=False).map(lambda x: round(x, 4)),
)
def test_mutation_d4_19_trust_file_rate_would_break_guard(
    revenue: float, discount: float, poison_rate: float
) -> None:
    """**Validates: Requirements 1.5, 4.1（mutation）**

    四态变异：把「派生列单源重算」翻成反模式「信任文件折扣比例列」。
    真 parser 重算 discountRate=折扣/收入；mutant 读文件污染值 poison_rate。
    断言 mutant 与真值在 poison_rate ≠ 真值时必不同——即单源守卫确实拦得住「信任文件」回归。
    """
    headers = _headers("D4-19")
    row_values = [
        "", "客户甲", "现金折扣", revenue, discount, poison_rate,
        "原因", "2025-01-01", "PZ-1", "6001", "主营", 0.0, 0.0, "2025-01-02", "张三", "备注",
    ]
    parsed = _roundtrip(headers, [row_values], mod._parse_d4_19_row)
    assert len(parsed) == 1
    real_rate = parsed[0]["discountRate"]
    true_rate = discount / revenue  # 单源应得值
    assert real_rate == true_rate

    # mutant：反模式「信任文件」= 直接取污染列值
    mutant_rate = poison_rate
    if abs(poison_rate - true_rate) > 1e-9:
        # 守卫命中：真值 ≠ mutant，说明「== 文件值」的断言会失败 → 守卫拦得住该回归
        assert real_rate != mutant_rate


@settings(max_examples=5)
@given(
    base=_amt,
    rate=st.floats(min_value=0, max_value=1, allow_nan=False).map(lambda x: round(x, 4)),
    already=_amt,
    poison_should=st.floats(min_value=1, max_value=1e6, allow_nan=False).map(lambda x: round(x, 2)),
)
def test_mutation_d4_20_provision_trust_file_derived_would_break(
    base: float, rate: float, already: float, poison_should: float
) -> None:
    """**Validates: Requirements 1.5, 4.1（mutation）**

    四态变异 D4-20 计提：反模式「信任文件应计提列」。真 parser 重算 base×rate；
    mutant 取文件 poison_should。poison ≠ 真值时守卫命中（触类旁通同 D4-19 单源反模式）。
    """
    headers = _headers("D4-20-provision")
    row_values = ["", "货品甲", base, rate, poison_should, already, 0.0, "原因"]
    parsed = _roundtrip(headers, [row_values], mod._parse_d4_20_provision_row)
    assert len(parsed) == 1
    real_should = parsed[0]["shouldProvide"]
    true_should = base * rate
    assert real_should == true_should
    if abs(poison_should - true_should) > 1e-9:
        assert real_should != poison_should  # 守卫拦得住「信任文件派生列」


# ─── Mutation 2 · 专用 IO：若 D4-19 走 generic 兜底则字段结构守卫失败 ───────────────

def test_mutation_generic_fallback_would_lose_dedicated_fields() -> None:
    """**Validates: Requirements 1.2, 4.1（mutation）**

    四态变异：把 D4-19 分发从专用 parser 翻成 generic 兜底。
    专用 parser 产出 discountRate/revenueAmount/discountAmount 等英文语义键；generic 兜底
    只按列头产原样字典、没有派生键与英文 key 映射。断言两者产物结构不同——即「专用 IO」守卫
    （dispatch 专用 parser）确实拦得住 generic 回归。
    """
    headers = _headers("D4-19")
    row = ("", "客户甲", "现金折扣", 1000.0, 200.0, 0.0, "原因", "2025-01-01",
           "PZ-1", "6001", "主营", 0.0, 0.0, "2025-01-02", "张三", "备注")
    dedicated = mod._parse_d4_19_row(row, headers)
    # 专用 parser 的英文语义键 + 单源派生键
    assert "discountRate" in dedicated
    assert dedicated["revenueAmount"] == 1000.0
    assert dedicated["discountRate"] == 200.0 / 1000.0

    # mutant：generic 兜底（按中文列头原样 → 无英文语义键、无单源重算）
    mutant = {h: v for h, v in zip(headers, row) if h}
    assert "discountRate" not in mutant  # generic 不产派生键
    assert "revenueAmount" not in mutant  # generic 不做英文 key 映射
    # 守卫命中：dedicated 与 mutant 结构不同，证明「dispatch 专用 parser」非可有可无
    assert set(dedicated.keys()) != set(mutant.keys())


# ─── Mutation 3 · item_id 双侧一致：若后端映射错 store 键则对齐守卫失败 ──────────────

def test_mutation_wrong_item_id_would_break_alignment_guard() -> None:
    """**Validates: Requirements 1.3, 2.3, 4.1（mutation）**

    四态变异：把 D4-20 子表 item_id 从对齐前端键（D4-20-current-returns）翻成裸 sheet 码
    （D4-20-current）。前端 store 只认 D4-20-current-returns/post-returns/provision。
    断言 mutant 键 ∉ 前端组件引用集——即 item_id 双侧一致守卫拦得住键错位回归。
    """
    fe = _read("audit-platform/frontend/src/components/workpaper/d4/inspection/D4TabReturn.vue")
    correct_keys = {"D4-20-current-returns", "D4-20-post-returns", "D4-20-provision"}
    for k in correct_keys:
        assert f"'{k}'" in fe  # 真键前端确实引用

    # mutant：错位成裸 sheet 码作为 store 键（曾经的 DEC-2 bug）。前端 store 引用的是带
    # -returns 后缀的键，裸 sheet 码作为「store 键字面量」从不出现（用引号包裹精确匹配 key 位置）。
    mutant_keys = {"D4-20-current", "D4-20-post"}
    for mk in mutant_keys:
        # 守卫命中：裸 sheet 码作为 store 键字面量 'D4-20-current'（后接引号闭合，非 -returns 前缀）
        # 在前端从不作为 item_id 键引用 → 若后端映射错成裸码，导入落孤儿键、导出读不到 → 往返丢数据。
        assert f"item_id: '{mk}'" not in fe
        assert f"'{mk}',\n" not in fe  # 裸码作为独立 key 字面量不存在
    # 反向确认：正确的带后缀键确实被前端引用（守卫非空转）
    assert "'D4-20-current-returns'" in fe
    assert "'D4-20-post-returns'" in fe

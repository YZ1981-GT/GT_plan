"""GS2 —— 16 张 X-3 调整分录汇总表「列面对源模板」守卫（openpyxl 直读）

spec: ``x3-adjustment-entry-import-export``，Wave 0 任务 1.2（判据先行）
Requirements: 3.1 / 3.2 / 3.6 / 3.7 / 3.8 / 11.1

## 判据形态：真实读文件，不是「源码里有这 10 个字面量」

R3.7 明文要求「比对列头时 SHALL 用 openpyxl 直读源模板取值，源码字符串 SHALL
仅作为被比对的一方」。所以本文件**不写任何列头字面量当期望值** —— 期望值每次
运行都从 ``backend/wp_templates/`` 对应 tab 的第 5 行现读，被比对的一方才是源码
（``COLUMN_ORDER``）与清单（``adjustment_ie_contract.json`` 的 ``column_map``）。

平台已多次踩过「grep 式守卫只查字符串存在」这一类假绿：源码里确实有那 10 个字，
守卫就绿了，而实际导出的列序/列数早就跑偏。此处一律以 workbook 实读值为准。

## 类 A / 类 B 分写（防「文件没找到导致断言空转」）

- **类 A（当前应全绿）**：catalog 16 条在位 · 16 个 xlsx 找得到 · tab 名逐字命中 ·
  第 5 行 10 个标签且 16 张逐字一致 · 第 5 行确实是表头行 · 比较器对替身打红 ·
  工厂 7 列 ``_ADJ_HEADERS`` 不构成列面真源 · 模板目录读取前后逐条未变。
  **类 A 全绿才说明「真读到了东西」**；类 A 一红，类 B 的红就不可解读。
- **类 B（当前应全红）**：与派生常量 ``COLUMN_ORDER``（Wave 2 任务 4.1 交付）和
  清单 ``column_map``（Wave 1 任务 2.1 交付）的三向比对。红是本任务的预期产出，
  消息里标明「尚未实现/尚未登记」+ 承接任务号。

## 三向比对（R3.1 / R3.2 / R3.6）

    清单 column_map[].col  ↔  派生常量 COLUMN_ORDER  ↔  openpyxl 实读第 5 行

三向逐字相等（等势 + 同序）才算收口。任一向缺失即红，且红的消息指出缺哪一向。

## ``backend/wp_templates/`` 只读（R11.1）

全链路只经 ``load_workbook(..., read_only=True)``；读取前后对整个目录做
``(相对路径, size, mtime_ns)`` 快照并断言逐条未变。``~$`` 锁文件跳过 —— 用户可能
正开着 WPS，锁文件出现/消失不属于「模板被改」。
"""

from __future__ import annotations

import ast
import importlib
import json
import re
from collections.abc import Sequence
from pathlib import Path

import pytest
from openpyxl import load_workbook

_REPO = Path(__file__).resolve().parents[2]
_TEMPLATES_DIR = _REPO / "backend" / "wp_templates"
_CATALOG = _REPO / "backend" / "data" / "acnr" / "global_catalog.json"
_CONTRACT = _REPO / "backend" / "data" / "adjustment_ie_contract.json"
_EVIDENCE_DUMP = (
    _REPO
    / ".kiro"
    / "specs"
    / "x3-adjustment-entry-import-export"
    / "evidence"
    / "source_template_x3_headers.txt"
)

#: Wave 2 任务 4.1 的唯一真源出口模块（当前尚不存在 ⇒ 类 B 全红）
_IMPL_MODULE = "app.routers.wp_render_strategies._x3_adjustment_import_export"

#: 源模板列头所在行。5 不是拍脑袋 —— ``test_class_a_header_row_is_row_5`` 用
#: 「16 张逐字一致 且 恰 10 个非空标签」这个结构判据反证第 1~8 行里只有它满足。
_HEADER_ROW = 5

#: 每张 X-3 的列数（源模板实测）。仅用于「读到的东西是不是那张表」的自检，
#: 期望列**内容**一律现读，不在此写字面量。
_EXPECTED_COL_COUNT = 10

#: 本 spec 作业面 = 16 张 X-3。Wave 0 阶段清单尚未迁入（15 张在 ``exempt``、
#: ``L6-3`` 连登记都没有），故作业面只能先在守卫里声明；
#: ``test_class_b_worksheet_scope_locked_by_contract`` 负责在 Wave 1 之后把它与
#: 清单 ``sheets`` 段双向锁死，防止这里变成第二份真源。
_X3_SHEET_CODES: tuple[str, ...] = (
    "L2-3",
    "L6-3",
    "M1-3",
    "M2-3",
    "M3-3",
    "M4-3",
    "M5-3",
    "M6-3",
    "M7-3",
    "M8-3",
    "M9-3",
    "M10-3",
    "N1-3",
    "N2-3",
    "N3-3",
    "N5-3",
)

_PROBE_ROWS = 8  # 读前 8 行供「第 5 行是表头行」的结构反证用


# ═══════════════════════════════════════════════════════════════════════════
# 比较器（判据核心，纯函数 ⇒ 可用替身做反向自检）
# ═══════════════════════════════════════════════════════════════════════════


def diff_columns(expected: Sequence[str], candidate: Sequence[str]) -> list[str]:
    """把 candidate 与源模板第 5 行 expected 逐字比对，返回违规描述。

    返回空列表 == 等势且同序且逐字相等。检出五类违规：
    多列（R3.6）/ 缺列（R3.2 等势）/ 重复列 / 列数不等 / 列序不一致。
    """
    exp = list(expected)
    cand = list(candidate)
    problems: list[str] = []

    extra = [c for c in cand if c not in exp]
    if extra:
        problems.append(
            f"出现源模板第 {_HEADER_ROW} 行不存在的列（R3.6）：{extra}"
        )
    missing = [e for e in exp if e not in cand]
    if missing:
        problems.append(f"缺列（R3.2 列集须与第 {_HEADER_ROW} 行等势）：{missing}")
    dup = sorted({c for c in cand if cand.count(c) > 1})
    if dup:
        problems.append(f"重复列：{dup}")
    if len(cand) != len(exp):
        problems.append(f"列数不等：期望 {len(exp)}、实测 {len(cand)}")
    if not problems and cand != exp:
        problems.append(
            "列序不一致："
            + " / ".join(
                f"#{i + 1} 期望 {e!r} 实测 {c!r}"
                for i, (e, c) in enumerate(zip(exp, cand))
                if e != c
            )
        )
    return problems


# ═══════════════════════════════════════════════════════════════════════════
# 读取辅助
# ═══════════════════════════════════════════════════════════════════════════


def _snapshot_templates_dir() -> dict[str, tuple[int, int]]:
    """``backend/wp_templates/`` 的 (相对路径 → (size, mtime_ns)) 快照。

    跳过 ``~$`` 锁文件：用户开着 WPS 时它会出现/消失，与「模板被改」无关。
    """
    snap: dict[str, tuple[int, int]] = {}
    for path in _TEMPLATES_DIR.rglob("*"):
        if not path.is_file():
            continue
        if path.name.startswith("~$"):
            continue
        st = path.stat()
        snap[path.relative_to(_TEMPLATES_DIR).as_posix()] = (st.st_size, st.st_mtime_ns)
    return snap


def _norm_row(row: tuple[object, ...]) -> tuple[str, ...]:
    """单元格值 → 字符串元组，去掉尾部空列（中间空列保留，`……` 占位列不算空）。"""
    out = ["" if v is None else str(v) for v in row]
    while out and out[-1] == "":
        out.pop()
    return tuple(out)


def _read_probe_rows(path: Path, tab: str) -> dict[int, tuple[str, ...]]:
    """只读打开 workbook，取前 ``_PROBE_ROWS`` 行。"""
    wb = load_workbook(path, read_only=True, data_only=True)
    try:
        assert tab in wb.sheetnames, (
            f"{path.name} 里没有 tab {tab!r}（实有 {wb.sheetnames}）—— "
            "catalog sheet_name 与源模板 tab 名不再逐字一致（R3.8）"
        )
        ws = wb[tab]
        rows: dict[int, tuple[str, ...]] = {}
        for idx, row in enumerate(
            ws.iter_rows(min_row=1, max_row=_PROBE_ROWS, values_only=True), start=1
        ):
            rows[idx] = _norm_row(row)
        return rows
    finally:
        wb.close()


def _load_contract() -> dict:
    assert _CONTRACT.is_file(), f"契约清单缺失：{_CONTRACT}（守卫必须打红而非跳过）"
    return json.loads(_CONTRACT.read_text(encoding="utf-8"))


def _contract_column_map(sheet_code: str) -> list[str] | None:
    """清单 ``sheets[sheet_code].column_map[].col``；未登记返回 None。"""
    entry = _load_contract().get("sheets", {}).get(sheet_code)
    if not isinstance(entry, dict):
        return None
    cmap = entry.get("column_map")
    if not isinstance(cmap, list) or not cmap:
        return None
    return [str(item.get("col")) for item in cmap if isinstance(item, dict)]


def _import_impl():
    """导入 Wave 2 的实现模块；模块本身缺失返回 None。

    只吞「就是这个模块不存在」这一种 ModuleNotFoundError —— 模块存在但其内部
    import 写错时必须原样抛出，否则会把接线错误伪装成「尚未实现」（fail-open）。
    """
    try:
        return importlib.import_module(_IMPL_MODULE)
    except ModuleNotFoundError as exc:
        if exc.name == _IMPL_MODULE:
            return None
        raise


def _factory_headers(cycle: str) -> list[str] | None:
    """工厂模块 ``_{cycle}_import_export._ADJ_HEADERS``（7 列），缺失返回 None。"""
    name = f"app.routers.wp_render_strategies._{cycle.lower()}_import_export"
    try:
        mod = importlib.import_module(name)
    except ModuleNotFoundError:
        return None
    headers = getattr(mod, "_ADJ_HEADERS", None)
    return list(headers) if headers else None


# ═══════════════════════════════════════════════════════════════════════════
# Fixtures（顺序：先快照 → 再读 workbook → 最后比对快照）
# ═══════════════════════════════════════════════════════════════════════════


@pytest.fixture(scope="module")
def templates_snapshot_before() -> dict[str, tuple[int, int]]:
    """读 workbook 之前的模板目录快照（R11.1 只读自检的基线）。"""
    assert _TEMPLATES_DIR.is_dir(), (
        f"源模板目录缺失：{_TEMPLATES_DIR} —— 它是运行时权威模板库，不可缺"
    )
    snap = _snapshot_templates_dir()
    # 反向自检：规模异常小 ⇒ 快照函数失效，后面「未变」的结论不可信
    assert len(snap) >= 300, (
        f"模板目录只快照到 {len(snap)} 个文件（实测基线 400+）—— "
        "快照函数或目录本身异常，此时『未被修改』的结论不可信"
    )
    return snap


@pytest.fixture(scope="module")
def catalog_x3() -> dict[str, dict]:
    """catalog 里 16 张 X-3 的条目（sheet_code → 条目）。"""
    assert _CATALOG.is_file(), f"ACNR catalog 缺失：{_CATALOG}"
    data = json.loads(_CATALOG.read_text(encoding="utf-8"))
    sheets = data.get("sheets")
    assert isinstance(sheets, list) and len(sheets) > 1000, (
        f"catalog sheets 段异常（type={type(sheets).__name__}，"
        f"len={len(sheets) if isinstance(sheets, list) else 'n/a'}）—— 解析失效"
    )
    by_code: dict[str, dict] = {}
    for item in sheets:
        code = item.get("sheet_code")
        if code in _X3_SHEET_CODES:
            # 同 sheet_code 多条目会让「取哪条」变得不确定，直接打红
            assert code not in by_code, f"catalog 里 {code} 有多条条目，登记面不唯一"
            by_code[code] = item
    return by_code


@pytest.fixture(scope="module")
def template_paths(catalog_x3: dict[str, dict]) -> dict[str, Path]:
    """16 张 X-3 → 其源模板 xlsx 路径（走平台自己的 finder，不写文件名字面量）。"""
    from app.services.wp_template_finder import find_template_file

    paths: dict[str, Path] = {}
    for code in _X3_SHEET_CODES:
        entry = catalog_x3.get(code) or {}
        parent = entry.get("parent_wp_code") or code.split("-")[0]
        found = find_template_file(parent)
        assert found is not None, (
            f"{code}：find_template_file({parent!r}) 找不到源模板 xlsx —— "
            "列面判据无从取值，后面所有比对都会空转"
        )
        assert _TEMPLATES_DIR in found.parents, (
            f"{code}：定位到 {found}，不在 {_TEMPLATES_DIR} 下 —— "
            "R3.1 要求列头取自 backend/wp_templates/（参考副本已落后，不可作真源）"
        )
        paths[code] = found
    return paths


@pytest.fixture(scope="module")
def probe_rows(
    templates_snapshot_before: dict[str, tuple[int, int]],
    catalog_x3: dict[str, dict],
    template_paths: dict[str, Path],
) -> dict[str, dict[int, tuple[str, ...]]]:
    """16 张 X-3 tab 的前 8 行实读值（唯一的真实取值处）。"""
    assert templates_snapshot_before, "只读基线快照为空 —— fixture 顺序被破坏"
    out: dict[str, dict[int, tuple[str, ...]]] = {}
    for code in _X3_SHEET_CODES:
        tab = (catalog_x3.get(code) or {}).get("sheet_name")
        assert tab, f"{code}：catalog 无 sheet_name，无法定位 tab（R3.8）"
        out[code] = _read_probe_rows(template_paths[code], str(tab))
    assert len(out) == len(_X3_SHEET_CODES), "实读覆盖面不足 16 张 —— 断言会空转"
    return out


@pytest.fixture(scope="module")
def row5_labels(
    probe_rows: dict[str, dict[int, tuple[str, ...]]],
) -> dict[str, tuple[str, ...]]:
    """16 张 X-3 第 5 行的 10 个标签（openpyxl 实读，三向比对的第一向）。"""
    return {code: rows[_HEADER_ROW] for code, rows in probe_rows.items()}


@pytest.fixture(scope="module")
def expected_columns(row5_labels: dict[str, tuple[str, ...]]) -> tuple[str, ...]:
    """16 张共同的列面（实读一致才有单一期望值，否则由类 A 打红）。"""
    variants = set(row5_labels.values())
    assert len(variants) == 1, (
        "16 张 X-3 第 5 行不再逐字一致，不存在单一列面期望值：\n"
        + "\n".join(f"  {c}: {v}" for c, v in sorted(row5_labels.items()))
    )
    return next(iter(variants))


# ═══════════════════════════════════════════════════════════════════════════
# 类 A —— 独立口径，当前应全绿（证明「真读到了东西」）
# ═══════════════════════════════════════════════════════════════════════════


def test_class_a_catalog_has_16_x3_entries(catalog_x3: dict[str, dict]) -> None:
    """catalog 16 条 X-3 在位，且 sheet_name 非空、class_code 为调整分录。"""
    missing = [c for c in _X3_SHEET_CODES if c not in catalog_x3]
    assert not missing, f"catalog 里找不到这些 X-3：{missing} —— 作业面定义与 catalog 脱节"
    bad: list[str] = []
    for code in _X3_SHEET_CODES:
        entry = catalog_x3[code]
        if not entry.get("sheet_name"):
            bad.append(f"{code}: sheet_name 为空")
        if entry.get("class_code") != "F-调整分录":
            bad.append(f"{code}: class_code={entry.get('class_code')!r}")
    assert not bad, "catalog 条目字段异常：\n" + "\n".join(f"  {b}" for b in bad)


def test_class_a_all_16_templates_resolved(template_paths: dict[str, Path]) -> None:
    """16 个源模板 xlsx 全部找得到（否则后面的列面断言全是空转）。"""
    assert len(template_paths) == len(_X3_SHEET_CODES)
    for code, path in sorted(template_paths.items()):
        assert path.is_file(), f"{code}: {path} 不是文件"
        assert path.suffix.lower() in (".xlsx", ".xlsm"), f"{code}: {path.name} 非 xlsx"


def test_class_a_catalog_sheet_name_matches_workbook_tab(
    catalog_x3: dict[str, dict],
    template_paths: dict[str, Path],
    probe_rows: dict[str, dict[int, tuple[str, ...]]],
) -> None:
    """R3.8：catalog ``sheet_name`` 与源模板 tab 名逐字一致（N 族多「表」字、L2 带业务前缀）。

    ``probe_rows`` 已经用该名字成功打开过 tab（``_read_probe_rows`` 内部断言），
    这里再把「逐字」这一层单独钉死：不许靠 strip/模糊匹配蒙对。
    """
    offenders: list[str] = []
    for code in _X3_SHEET_CODES:
        name = str((catalog_x3[code] or {}).get("sheet_name"))
        wb = load_workbook(template_paths[code], read_only=True, data_only=True)
        try:
            tabs = list(wb.sheetnames)
        finally:
            wb.close()
        if name not in tabs:
            near = [t for t in tabs if code in t]
            offenders.append(f"{code}: catalog={name!r} 不在 tab 列表；形近项={near}")
    assert not offenders, (
        "catalog sheet_name 与源模板 tab 名不再逐字一致（R3.8）：\n"
        + "\n".join(f"  {o}" for o in offenders)
    )
    assert probe_rows, "probe_rows 为空 —— 实读未发生"


def test_class_a_row5_has_10_labels_and_is_identical(
    row5_labels: dict[str, tuple[str, ...]],
) -> None:
    """16 张第 5 行各 10 个非空标签、无首尾空白、且 16 张逐字一致。"""
    problems: list[str] = []
    for code, labels in sorted(row5_labels.items()):
        nonempty = [x for x in labels if x != ""]
        if len(labels) != _EXPECTED_COL_COUNT or len(nonempty) != _EXPECTED_COL_COUNT:
            problems.append(
                f"{code}: 第 {_HEADER_ROW} 行 {len(labels)} 列 / "
                f"{len(nonempty)} 个非空（期望各 {_EXPECTED_COL_COUNT}）→ {labels}"
            )
        for x in labels:
            if x != x.strip():
                problems.append(f"{code}: 标签 {x!r} 带首尾空白，逐字比对会假红")
    assert not problems, "第 5 行结构异常：\n" + "\n".join(f"  {p}" for p in problems)

    variants = sorted(set(row5_labels.values()))
    assert len(variants) == 1, (
        f"16 张第 {_HEADER_ROW} 行出现 {len(variants)} 种列面（实证基线：逐字一致）：\n"
        + "\n".join(f"  {v}" for v in variants)
    )


def test_class_a_header_row_is_row_5(
    probe_rows: dict[str, dict[int, tuple[str, ...]]],
) -> None:
    """反证第 5 行就是表头行 —— 不把「第 5 行」当无据假设。

    判据：第 1~8 行里，同时满足「16 张逐字一致」与「恰 10 个非空标签」的行**只有**
    第 5 行（第 1 行虽一致但只有 1 个标签，第 2~4 行逐张不同，第 6 行 L2 有示例数据）。
    """
    qualifying: list[int] = []
    for row_no in range(1, _PROBE_ROWS + 1):
        values = {rows.get(row_no, ()) for rows in probe_rows.values()}
        if len(values) != 1:
            continue
        only = next(iter(values))
        if len([x for x in only if x != ""]) == _EXPECTED_COL_COUNT:
            qualifying.append(row_no)
    assert qualifying == [_HEADER_ROW], (
        f"「16 张一致且恰 {_EXPECTED_COL_COUNT} 个标签」的行为 {qualifying}，"
        f"不是 [{_HEADER_ROW}] —— 表头行位置变了，_HEADER_ROW 须重新核定"
    )


def test_class_a_live_read_matches_evidence_snapshot(
    row5_labels: dict[str, tuple[str, ...]],
) -> None:
    """实读值与 spec 留存的实证快照逐字相等（快照是证据，实读才是真源）。

    双向价值：模板被改 → 这里红；快照过期/被谁重写坏了 → 也红。
    """
    assert _EVIDENCE_DUMP.is_file(), f"实证快照缺失：{_EVIDENCE_DUMP}"
    text = _EVIDENCE_DUMP.read_text(encoding="utf-8")
    dumped = [ast.literal_eval(m) for m in re.findall(r"^\s*r5:\s*(\[.*\])\s*$", text, re.M)]
    assert len(dumped) == len(_X3_SHEET_CODES), (
        f"快照里解析出 {len(dumped)} 条 r5 行（期望 {len(_X3_SHEET_CODES)}）—— "
        "解析锚点未命中或快照格式变了（ANCHOR-MISS，不是被测对象的问题）"
    )
    snapshot_variants = {tuple(str(x) for x in row) for row in dumped}
    assert len(snapshot_variants) == 1, f"快照内部就不一致：{sorted(snapshot_variants)}"
    snap = next(iter(snapshot_variants))
    for code, labels in sorted(row5_labels.items()):
        assert labels == snap, (
            f"{code} 实读第 {_HEADER_ROW} 行与实证快照不符：\n"
            f"  实读={labels}\n  快照={snap}\n"
            "→ 要么 backend/wp_templates/ 被改（违反 R11.1），要么快照已过期需重生成"
        )


def test_class_a_factory_adj_headers_is_not_column_source(
    expected_columns: tuple[str, ...],
) -> None:
    """R3.2：工厂 ``_ADJ_HEADERS``（7 列）与第 5 行不等势 ⇒ 不构成列面真源。

    这是一条独立事实断言（不依赖尚未存在的实现）：16 个工厂模块的列面都缺列，
    连「类别」一列的标签都与源模板不逐字相同。谁把它当列面真源，导出就少列。
    """
    checked = 0
    for code in _X3_SHEET_CODES:
        cycle = code.split("-")[0]
        headers = _factory_headers(cycle)
        assert headers, f"{cycle} 工厂模块的 _ADJ_HEADERS 取不到 —— 本断言已空转"
        checked += 1
        problems = diff_columns(expected_columns, headers)
        assert problems, (
            f"{cycle} 工厂 _ADJ_HEADERS 竟与源模板第 {_HEADER_ROW} 行等势且同序 —— "
            "与 R3.2 的实证前提冲突，须重新核定列面真源"
        )
        assert len(headers) < len(expected_columns), (
            f"{cycle} 工厂列数 {len(headers)} 不小于源模板 {len(expected_columns)}，"
            "R3.2「工厂 7 列缺 3 列」的前提已变"
        )
    assert checked == len(_X3_SHEET_CODES), f"只核了 {checked} 个工厂模块"


def test_class_a_templates_dir_unchanged_after_reads(
    templates_snapshot_before: dict[str, tuple[int, int]],
    probe_rows: dict[str, dict[int, tuple[str, ...]]],
    row5_labels: dict[str, tuple[str, ...]],
) -> None:
    """R11.1：读取前后 ``backend/wp_templates/`` 逐条 (size, mtime_ns) 未变。

    ``probe_rows`` / ``row5_labels`` 作为入参，保证 16 次 workbook 打开已经发生过
    （否则这条自检等于「什么都没读也说没改」）。
    """
    assert len(row5_labels) == len(_X3_SHEET_CODES), "实读未覆盖 16 张，自检不成立"
    assert probe_rows, "实读未发生"

    after = _snapshot_templates_dir()
    changed = {
        rel: (templates_snapshot_before[rel], after[rel])
        for rel in templates_snapshot_before.keys() & after.keys()
        if templates_snapshot_before[rel] != after[rel]
    }
    removed = sorted(templates_snapshot_before.keys() - after.keys())
    added = sorted(after.keys() - templates_snapshot_before.keys())
    assert not (changed or removed or added), (
        "读取源模板后 backend/wp_templates/ 发生变化（违反 R11.1；"
        "也可能是并发会话/WPS 正在写该目录）：\n"
        + "".join(f"  改动 {r}: {b} → {a}\n" for r, (b, a) in sorted(changed.items()))
        + "".join(f"  消失 {r}\n" for r in removed)
        + "".join(f"  新增 {r}\n" for r in added)
    )


# ── 比较器反向自检（替身：少一列 / 顺序错 / 多一列 / 标签不逐字 / 重复列）──


def test_class_a_diff_columns_accepts_exact_match(
    expected_columns: tuple[str, ...],
) -> None:
    assert diff_columns(expected_columns, list(expected_columns)) == []


def test_class_a_diff_columns_rejects_missing_column(
    expected_columns: tuple[str, ...],
) -> None:
    """替身：少一列 ⇒ 必须打红。"""
    stand_in = list(expected_columns)[:-1]
    problems = diff_columns(expected_columns, stand_in)
    assert problems, "少一列的替身竟被判通过 —— 比较器有缺陷（等势判据空转）"
    assert any("缺列" in p for p in problems), problems


def test_class_a_diff_columns_rejects_wrong_order(
    expected_columns: tuple[str, ...],
) -> None:
    """替身：等势但顺序错 ⇒ 必须打红（导出列序错，用户看到的是错位数据）。"""
    stand_in = list(expected_columns)
    stand_in[0], stand_in[1] = stand_in[1], stand_in[0]
    problems = diff_columns(expected_columns, stand_in)
    assert problems, "顺序错的替身竟被判通过 —— 比较器只比了集合没比序"
    assert any("列序不一致" in p for p in problems), problems


def test_class_a_diff_columns_rejects_extra_column(
    expected_columns: tuple[str, ...],
) -> None:
    """替身：多一列（源模板第 5 行不存在的列）⇒ 必须打红（R3.6 的正面判据）。"""
    stand_in = [*expected_columns, "序号"]
    problems = diff_columns(expected_columns, stand_in)
    assert problems, "多一列的替身竟被判通过 —— R3.6 判据空转"
    assert any(f"第 {_HEADER_ROW} 行不存在的列" in p for p in problems), problems


def test_class_a_diff_columns_rejects_non_literal_label(
    expected_columns: tuple[str, ...],
) -> None:
    """替身：标签不逐字（工厂的「类别」vs 源模板「类别（报表调整/账项调整/其他）」）⇒ 必须打红。"""
    long_label = next(
        (x for x in expected_columns if len(x) > 4 and "（" in x), None
    )
    assert long_label, f"取不到带括号的长标签做替身：{expected_columns}"
    stand_in = [long_label.split("（")[0] if x == long_label else x for x in expected_columns]
    problems = diff_columns(expected_columns, stand_in)
    assert problems, "标签被截短的替身竟被判通过 —— 比对不是逐字的"
    assert any("不存在的列" in p for p in problems) and any(
        "缺列" in p for p in problems
    ), problems


def test_class_a_diff_columns_rejects_duplicate_column(
    expected_columns: tuple[str, ...],
) -> None:
    """替身：重复列 ⇒ 必须打红。"""
    stand_in = list(expected_columns)
    stand_in[-1] = stand_in[0]
    problems = diff_columns(expected_columns, stand_in)
    assert problems, "重复列的替身竟被判通过"
    assert any("重复列" in p for p in problems), problems


# ═══════════════════════════════════════════════════════════════════════════
# 类 B —— 被测实现，当前应全红（红是本任务的预期产出）
#
# 红的两个承接任务：
#   · COLUMN_ORDER / X3_SHEET_SPECS  → 尚未实现（Wave 2 任务 4.1）
#   · 清单 column_map / sheets 迁入   → 尚未登记（Wave 1 任务 2.1）
# ═══════════════════════════════════════════════════════════════════════════


def test_class_b_column_order_is_importable() -> None:
    """派生常量 ``COLUMN_ORDER`` 存在且非空（三向比对的第二向）。"""
    mod = _import_impl()
    assert mod is not None, (
        f"{_IMPL_MODULE} 不存在 —— 尚未实现（Wave 2 任务 4.1）。"
        "该模块是 COLUMN_ORDER 的唯一定义处（由清单 column_map[].col 派生）"
    )
    order = getattr(mod, "COLUMN_ORDER", None)
    assert order, (
        f"{_IMPL_MODULE} 里没有 COLUMN_ORDER —— 尚未实现（Wave 2 任务 4.1）"
    )
    assert len(order) == len(set(order)), f"COLUMN_ORDER 有重复列：{order}"


def test_class_b_contract_column_map_registered() -> None:
    """清单 ``sheets[X-3].column_map`` 16 张齐备（三向比对的第三向）。"""
    contract = _load_contract()
    sheets = contract.get("sheets", {})
    missing_entry = [c for c in _X3_SHEET_CODES if c not in sheets]
    missing_map = [
        c for c in _X3_SHEET_CODES if c in sheets and not _contract_column_map(c)
    ]
    assert not (missing_entry or missing_map), (
        "契约清单尚未登记 column_map —— 尚未登记（Wave 1 任务 2.1）：\n"
        f"  未迁入 sheets 段：{missing_entry}\n"
        f"  已迁入但无 column_map：{missing_map}\n"
        "→ column_map[].col 是那 10 个列头字面量的唯一出现处，COLUMN_ORDER 由它派生"
    )


def test_class_b_worksheet_scope_locked_by_contract() -> None:
    """作业面 16 张与清单 ``sheets`` 段双向锁死（防守卫成为第二份真源）。"""
    contract = _load_contract()
    sheets = set(contract.get("sheets", {}))
    exempt = contract.get("exempt", {})
    still_exempt = [
        c
        for c in _X3_SHEET_CODES
        if isinstance(exempt.get(c), dict)
        and exempt[c].get("kind") == "no_backend_spec"
    ]
    not_registered = [c for c in _X3_SHEET_CODES if c not in sheets]
    assert not (not_registered or still_exempt), (
        "16 张 X-3 尚未从 exempt 迁入 sheets —— 尚未登记（Wave 1 任务 2.1）：\n"
        f"  未在 sheets 段：{not_registered}\n"
        f"  仍挂 exempt.no_backend_spec：{still_exempt}"
    )


def test_class_b_three_way_column_alignment(
    row5_labels: dict[str, tuple[str, ...]],
) -> None:
    """三向比对：清单 ``column_map[].col`` ↔ ``COLUMN_ORDER`` ↔ openpyxl 实读。

    实读值是期望方（R3.7），另两向都是被比对方。任一向缺失即红并指出缺哪一向。
    """
    mod = _import_impl()
    order = list(getattr(mod, "COLUMN_ORDER", []) or []) if mod else []
    unavailable: list[str] = []
    if not order:
        unavailable.append(
            f"第二向 COLUMN_ORDER 取不到（{_IMPL_MODULE}）—— 尚未实现（Wave 2 任务 4.1）"
        )

    failures: list[str] = []
    for code in _X3_SHEET_CODES:
        expected = list(row5_labels[code])
        if order:
            problems = diff_columns(expected, order)
            if problems:
                failures.append(f"{code} COLUMN_ORDER vs 实读：" + "；".join(problems))
        cmap = _contract_column_map(code)
        if cmap is None:
            unavailable.append(f"第三向 {code} 清单 column_map 未登记")
            continue
        problems = diff_columns(expected, cmap)
        if problems:
            failures.append(f"{code} 清单 column_map vs 实读：" + "；".join(problems))
        if order and cmap != order:
            failures.append(
                f"{code} 清单 column_map 与 COLUMN_ORDER 不一致（派生关系断了）："
                f"{cmap} != {order}"
            )

    assert not (unavailable or failures), (
        "三向比对未收口（尚未登记 Wave 1 任务 2.1 / 尚未实现 Wave 2 任务 4.1）：\n"
        + "".join(f"  [缺一向] {u}\n" for u in unavailable)
        + "".join(f"  [不一致] {f}\n" for f in failures)
    )


def test_class_b_column_order_has_no_column_absent_from_row5(
    expected_columns: tuple[str, ...],
) -> None:
    """R3.6 正面判据：``COLUMN_ORDER`` 出现源模板第 5 行不存在的列即红。"""
    mod = _import_impl()
    assert mod is not None, (
        f"{_IMPL_MODULE} 不存在 —— 尚未实现（Wave 2 任务 4.1），R3.6 判据无对象可查"
    )
    order = list(getattr(mod, "COLUMN_ORDER", []) or [])
    assert order, f"COLUMN_ORDER 为空 —— 尚未实现（Wave 2 任务 4.1）"
    extra = [c for c in order if c not in expected_columns]
    assert not extra, (
        f"COLUMN_ORDER 里有源模板第 {_HEADER_ROW} 行不存在的列（R3.6）：{extra}\n"
        f"  源模板第 {_HEADER_ROW} 行={expected_columns}"
    )


def test_class_b_impl_does_not_reuse_factory_headers(
    expected_columns: tuple[str, ...],
) -> None:
    """R3.2：实现的列面不得退回工厂 7 列 ``_ADJ_HEADERS``。

    主判据是行为（``COLUMN_ORDER`` 与工厂 7 列不等 且 与实读等势同序）；
    顺带查一眼实现模块有没有直接引用工厂常量（辅助判据，不单独成立）。
    """
    mod = _import_impl()
    assert mod is not None, (
        f"{_IMPL_MODULE} 不存在 —— 尚未实现（Wave 2 任务 4.1）"
    )
    order = list(getattr(mod, "COLUMN_ORDER", []) or [])
    assert order, "COLUMN_ORDER 为空 —— 尚未实现（Wave 2 任务 4.1）"

    factory = _factory_headers(_X3_SHEET_CODES[0].split("-")[0]) or []
    assert order != factory, (
        f"COLUMN_ORDER 与工厂 _ADJ_HEADERS 相同（{factory}）—— "
        "工厂 7 列缺 3 列，不构成列面真源（R3.2）"
    )
    assert diff_columns(expected_columns, order) == [], (
        "COLUMN_ORDER 与源模板第 5 行不等势/不同序："
        + "；".join(diff_columns(expected_columns, order))
    )

    src = Path(str(mod.__file__)).read_text(encoding="utf-8")
    assert "_ADJ_HEADERS" not in src, (
        f"{_IMPL_MODULE} 引用了工厂常量 _ADJ_HEADERS —— 列面真源须是清单 column_map"
    )


def test_class_b_sheet_name_taken_from_catalog(
    catalog_x3: dict[str, dict],
    probe_rows: dict[str, dict[int, tuple[str, ...]]],
) -> None:
    """R3.8：``X3_SHEET_SPECS[code].sheet_name`` == catalog ``sheet_name``（已与 tab 名逐字锁死）。"""
    mod = _import_impl()
    assert mod is not None, (
        f"{_IMPL_MODULE} 不存在 —— 尚未实现（Wave 2 任务 4.1）；"
        "导出 sheet 名须取 catalog sheet_name（N 族多「表」字、L2 带业务前缀）"
    )
    specs = getattr(mod, "X3_SHEET_SPECS", None)
    assert specs, f"{_IMPL_MODULE} 里没有 X3_SHEET_SPECS —— 尚未实现（Wave 2 任务 4.1）"

    assert probe_rows, "实读未发生，R3.8 的『与 tab 名逐字一致』这一层未被证明"
    offenders: list[str] = []
    for code in _X3_SHEET_CODES:
        spec = specs.get(code)
        if spec is None:
            offenders.append(f"{code}: X3_SHEET_SPECS 无该 sheet")
            continue
        declared = getattr(spec, "sheet_name", None)
        expected = (catalog_x3[code] or {}).get("sheet_name")
        if declared != expected:
            offenders.append(f"{code}: 实现={declared!r} catalog={expected!r}")
    assert not offenders, (
        "导出 sheet 名与 catalog sheet_name 不一致（R3.8）：\n"
        + "\n".join(f"  {o}" for o in offenders)
    )

"""N3 / N4 四表取数守卫（Property 3 / 4 / 5 / 8 / 13）+ 共享模块（Property 4/5/6）。

spec: `.kiro/specs/n345-four-table-extraction-alignment/`

拦下列已实证缺陷：

- **N3 原只查 `account_code == '2901'` 精确单行 + `.limit(1)`** → 客户按子科目挂账
  （活体 `2901.01/.02/.03` 普遍存在）时恒返回 0；且审定表 seed 把 TB 数据
  **全塞进「其他」一行**，分类信息全丢。
- **N3_SHEETS 曾列「附注披露信息」** → N3 源模板无该 sheet（披露与 N1 共节），
  宿主 `isHtmlSheet` 也认它 → 该 sheet 一旦出现即渲染空白 Tab（inert 残留）。
- **N4 曾 sum `tb_ledger` 的「Σ借 − Σ贷」** → 损益类含年末结转损益，该差结构性恒为 0。
- **N4 按 `6403` 子科目编码取税种** → 活体编码语义客户间冲突，只能按名称。
"""

from __future__ import annotations

import asyncio
import inspect
import re
import types
from pathlib import Path
from unittest.mock import AsyncMock

import pytest

from app.routers.wp_render_strategies import _n3_deferred_tax_liabilities as n3
from app.routers.wp_render_strategies import _n4_taxes_and_surcharges as n4
from app.services import deferred_tax_shared as shared

# 本文件位于 backend/tests/ → parents[1] = backend / parents[2] = 仓库根
BACKEND = Path(__file__).resolve().parents[1]
REPO_ROOT = Path(__file__).resolve().parents[2]


def _run(coro):
    return asyncio.run(coro)


class _Result:
    def __init__(self, rows=()):
        self._rows = list(rows)

    def fetchall(self):
        return self._rows

    def fetchone(self):
        return self._rows[0] if self._rows else None


def _bal(code, name="", opening=0.0, closing=0.0, debit=0.0, credit=0.0):
    return types.SimpleNamespace(
        account_code=code,
        account_name=name,
        opening_balance=opening,
        closing_balance=closing,
        debit_amount=debit,
        credit_amount=credit,
    )


def _tb(code, unadjusted=0.0):
    return types.SimpleNamespace(
        standard_account_code=code, unadjusted_amount=unadjusted, audited_amount=unadjusted
    )


def _ctx(result=None, *, raise_on_execute=False, results=None):
    db = AsyncMock()
    if raise_on_execute:
        db.execute = AsyncMock(side_effect=RuntimeError("boom"))
    elif results is not None:
        db.execute = AsyncMock(side_effect=list(results))
    else:
        db.execute = AsyncMock(return_value=result or _Result())
    return types.SimpleNamespace(db=db, project_id="p1", year=2025, wp_id="w1")


@pytest.fixture(autouse=True)
def _patch_active_filter(monkeypatch):
    async def _fake(db, table, project_id, year, **kw):  # noqa: ANN001
        return True

    for mod in (n3, n4):
        monkeypatch.setattr(mod, "get_active_filter", _fake)


# ═══════════════════════ 共享模块（Property 4/5/6）═══════════════════════


def test_shared_slots_order():
    assert shared.LIABILITY_SLOTS == [
        "depreciation",
        "afs_fv",
        "investment_property_fv",
        "lease",
        "other",
    ]


def test_shared_classify_total_and_priority():
    c = shared.classify_liability_subaccount
    assert c("投资性房地产公允价值变动") == "investment_property_fv"
    assert c("递延所得税负债_公允价值变动") == "afs_fv"
    assert c("递延所得税负债_固定资产加速折旧") == "depreciation"
    assert c("递延所得税负债_经营租赁相关") == "lease"
    assert c("使用权资产") == "lease"
    assert c("租赁形成") == "lease"
    for s in ("", None, "x", "！@#", "递延所得税负债"):
        assert c(s) in shared.LIABILITY_SLOTS


def test_shared_matches_n1_behaviour():
    """Property 6：共享模块与 N1 模块内实现输出等价（防迁移引入偏差）。"""
    from app.routers.wp_render_strategies import _n1_deferred_tax_assets as n1

    samples = [
        "投资性房地产公允价值变动", "递延所得税负债_公允价值变动",
        "递延所得税负债_固定资产加速折旧", "经营租赁相关", "使用权资产",
        "租赁形成", "购入摊销年限大于税法规定的资产", "其他", "", None,
    ]
    for s in samples:
        assert n1._classify_n1_liability_subaccount(s) == shared.classify_liability_subaccount(s)
    assert n1._N1_LIABILITY_SLOTS == shared.LIABILITY_SLOTS


def test_shared_leaf_rows_dotted_and_flat():
    assert {r.account_code for r in shared.leaf_rows([_bal("2901"), _bal("2901.01"), _bal("2901.02")])} == {
        "2901.01",
        "2901.02",
    }
    assert {r.account_code for r in shared.leaf_rows([_bal("2901"), _bal("290101")])} == {"290101"}


def test_shared_aggregate_skips_all_zero_and_abs():
    rows = [
        _bal("2901.01", "公允价值变动", 0.0, 0.0),
        _bal("2901.03", "经营租赁相关", -528013.88, -233512.19),
    ]
    out = shared.aggregate_by_slot(rows, shared.classify_liability_subaccount, absolute=True)
    assert "afs_fv" not in out
    assert out["lease"] == {"opening": 528013.88, "closing": 233512.19}


# ═══════════════════════ N3 ═══════════════════════


def test_n3_row_code_is_db_verified():
    assert n3._N3_ROW_CODE == "BS-067"


def test_n3_resolve_fail_open(monkeypatch):
    monkeypatch.setattr(
        n3, "resolve_report_line_account_codes", AsyncMock(side_effect=RuntimeError("x"))
    )
    assert _run(n3._resolve_account_codes(_ctx())) == ["2901"]


def test_n3_resolve_uses_report_config(monkeypatch):
    monkeypatch.setattr(
        n3, "resolve_report_line_account_codes", AsyncMock(return_value=["2901", "2902"])
    )
    assert _run(n3._resolve_account_codes(_ctx())) == ["2901", "2902"]


def test_n3_fetch_prefers_exact_account_row():
    rows = [
        _bal("2901", "递延所得税负债", -528013.88, -233512.19),
        _bal("2901.03", "经营租赁相关", -528013.88, -233512.19),
    ]
    out = _run(n3._fetch_tb_data(_ctx(_Result(rows))))
    assert out["begin_balance"] == 528013.88
    assert out["end_balance"] == 233512.19


def test_n3_fetch_falls_back_to_leaf_aggregation():
    """🔴 原实现 `.limit(1)` + 精确匹配：客户按子科目挂账时恒 0。"""
    rows = [
        _bal("2901.01", "公允价值变动", -1000.0, -1200.0),
        _bal("2901.03", "经营租赁相关", -2000.0, -2500.0),
    ]
    out = _run(n3._fetch_tb_data(_ctx(_Result(rows))))
    assert out["begin_balance"] == 3000.0
    assert out["end_balance"] == 3700.0


def test_n3_fetch_abs_normalizes_both_sign_conventions():
    """活体实测 2901 期末同时存在 -233512.19 与 200530.32 两种符号约定。"""
    a = _run(n3._fetch_tb_data(_ctx(_Result([_bal("2901", closing=-233512.19)]))))
    b = _run(n3._fetch_tb_data(_ctx(_Result([_bal("2901", closing=200530.32)]))))
    assert a["end_balance"] == 233512.19
    assert b["end_balance"] == 200530.32


def test_n3_fetch_query_error_fails_open():
    out = _run(n3._fetch_tb_data(_ctx(raise_on_execute=True)))
    assert out["end_balance"] == 0


def test_n3_prefill_by_slot_not_all_other():
    """🔴 原实现把 TB 数据全塞「其他」行；现按语义槽分类。"""
    rows = [
        _bal("2901.01", "递延所得税负债_公允价值变动", -1000.0, -1200.0),
        _bal("2901.02", "递延所得税负债_固定资产加速折旧", -2000.0, -2500.0),
        _bal("2901.03", "递延所得税负债_经营租赁相关", -528013.88, -233512.19),
    ]
    out = _run(n3._build_adjudication_prefill(_ctx(_Result(rows))))
    assert out["afs_fv"] == {"opening": 1000.0, "closing": 1200.0}
    assert out["depreciation"] == {"opening": 2000.0, "closing": 2500.0}
    assert out["lease"] == {"opening": 528013.88, "closing": 233512.19}
    assert "other" not in out


def test_n3_prefill_parent_only_returns_empty():
    out = _run(n3._build_adjudication_prefill(_ctx(_Result([_bal("2901", closing=-100.0)]))))
    assert out == {}


def test_n3_prefill_query_error_fails_open():
    out = _run(n3._build_adjudication_prefill(_ctx(raise_on_execute=True)))
    assert out == {}


# ─── Property 13: N3 无披露 inert 残留 ──────────────────────────────────────


def test_n3_sheets_has_no_disclosure_entry():
    """N3 源模板无「附注披露信息」sheet（披露与 N1 共节 五、30 / 八、31）。"""
    bad = [
        s["sheet_name"]
        for s in n3.N3_SHEETS
        if "附注" in s["sheet_name"] or "披露" in s["sheet_name"]
    ]
    assert bad == [], f"N3_SHEETS 不得含披露 sheet：{bad}"


def test_n3_sheets_matches_source_template():
    """交叉比对源模板 tab 名（openpyxl 直读），防再次凭空添 sheet。"""
    from openpyxl import load_workbook

    xlsx = BACKEND / "wp_templates" / "N" / "N3 递延所得税负债.xlsx"
    assert xlsx.exists(), f"源模板缺失：{xlsx}"
    real = load_workbook(xlsx, read_only=True).sheetnames
    assert not any("附注" in s for s in real), f"源模板竟有附注 sheet：{real}"
    # 声明的 sheet 数不应超过源模板（GT_Custom 不渲染）
    assert len(n3.N3_SHEETS) <= len(real)


# ═══════════════════════ N4 ═══════════════════════


def test_n4_row_code_is_db_verified():
    assert n4._N4_ROW_CODE == "IS-003"


def test_n4_resolve_fail_open(monkeypatch):
    monkeypatch.setattr(
        n4, "resolve_report_line_account_codes", AsyncMock(side_effect=RuntimeError("x"))
    )
    assert _run(n4._resolve_account_codes(_ctx())) == ["6403"]


def test_n4_period_amount_prefers_trial_balance():
    ctx = _ctx(results=[_Result([_tb("6403", 11_258_989.55)])])
    out = _run(n4._fetch_period_amount(ctx, ["6403"]))
    assert out["audited_amount"] == 11_258_989.55
    assert out["source"] == "trial_balance"


def test_n4_period_amount_falls_back_to_tb_balance_debit():
    ctx = _ctx(results=[_Result([]), _Result([_bal("6403", debit=3_231_580.11)])])
    out = _run(n4._fetch_period_amount(ctx, ["6403"]))
    assert out["audited_amount"] == 3_231_580.11
    assert out["source"] == "tb_balance"


def test_n4_no_ledger_debit_minus_credit():
    """损益类含年末结转损益 → 「Σ借 − Σ贷」结构性恒为 0（活体 9 项目全中）。"""
    src = re.sub(r"#[^\n]*", "", re.sub(r'"""[\s\S]*?"""', "", inspect.getsource(n4)))
    assert "TbLedger" not in src, "不得再从 tb_ledger sum 借贷（结转损益使其差恒为 0）"


def test_n4_classify_by_name_priority():
    c = n4._classify_n4_subaccount
    # 「地方教育附加」必须先于「教育费附加」（后者是前者子串）
    assert c("税金及附加_地方教育费附加") == "地方教育附加"
    assert c("税金及附加_教育费附加") == "教育费附加"
    assert c("税金及附加_城市维护建设税") == "城市维护建设税"
    assert c("城建税") == "城市维护建设税"
    assert c("车船税") == "车船税"
    assert c("税金及附加_车船使用税") == "车船税"
    assert c("印花税") == "印花税"
    assert c("税金及附加_土地使用税") == "城镇土地使用税"
    assert c("税金及附加_房产税") == "房产税"
    assert c("税金及附加_资源税") == "资源税"
    assert c("消费税") == "消费税"
    # 活体存在但无对应税种行 → 其他
    assert c("地方水利建设基金") == "其他"
    assert c("税金及附加_环境保护税") == "其他"
    for s in ("", None, "！@#"):
        assert c(s) in n4._N4_TAX_TYPES


def test_n4_tax_types_mirror_frontend():
    """后端税种清单必须与前端 `useN4Adjudication.DEFAULT_TAX_TYPES` 逐字同序。

    前端注释明确：命名与顺序须与 useN4Detail / useN4CrossSheet(N2 勾稽按名匹配) /
    附注三处一致 → 后端若漂移，预填就会有落不到行的孤儿税种。
    """
    fe = (
        REPO_ROOT
        / "audit-platform"
        / "frontend"
        / "src"
        / "components"
        / "workpaper"
        / "composables"
        / "useN4Adjudication.ts"
    )
    assert fe.exists(), f"前端文件缺失：{fe}"
    src = fe.read_text(encoding="utf-8")
    block = re.search(
        r"DEFAULT_TAX_TYPES:\s*Array<\{[^}]*\}>\s*=\s*\[(.*?)\]", src, re.S
    )
    assert block, "未能从前端源码抽出 DEFAULT_TAX_TYPES（正则失效？）"
    fe_types = re.findall(r"taxType:\s*'([^']+)'", block.group(1))
    assert fe_types, "抽取结果为空 → 正则失效"
    assert fe_types == n4._N4_TAX_TYPES, (
        f"后端税种清单与前端不一致：\n  后端 {n4._N4_TAX_TYPES}\n  前端 {fe_types}"
    )


def test_n4_prefill_aggregates_by_tax_type():
    rows = [
        _bal("6403", "税金及附加", debit=100.0),                 # 父级排除
        _bal("6403.02", "税金及附加_城市维护建设税", debit=1_022_169.44),
        _bal("6403.03", "税金及附加_教育费附加", debit=438_072.64),
        _bal("6403.04", "税金及附加_地方教育费附加", debit=292_048.39),
        _bal("6403.09", "税金及附加_印花税", debit=137_266.86),
        _bal("6403.99", "税金及附加_其他", debit=0.0),          # 全零跳过
    ]
    out = _run(n4._build_adjudication_prefill(_ctx(_Result(rows)), ["6403"]))
    assert out["城市维护建设税"] == 1_022_169.44
    assert out["教育费附加"] == 438_072.64
    assert out["地方教育附加"] == 292_048.39
    assert out["印花税"] == 137_266.86
    assert "其他" not in out
    # 父级 100 不得混入
    assert round(sum(out.values()), 2) == round(
        1_022_169.44 + 438_072.64 + 292_048.39 + 137_266.86, 2
    )


def test_n4_prefill_parent_only_returns_empty():
    out = _run(n4._build_adjudication_prefill(_ctx(_Result([_bal("6403", debit=100.0)])), ["6403"]))
    assert out == {}


def test_n4_prefill_query_error_fails_open():
    out = _run(n4._build_adjudication_prefill(_ctx(raise_on_execute=True), ["6403"]))
    assert out == {}

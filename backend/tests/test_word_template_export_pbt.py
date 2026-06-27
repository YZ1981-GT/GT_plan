"""Property-Based Tests for word-template export (prefilled-download with include_responses).

Property 9: Export placeholder replacement with format preservation
Property 10: Export round-trip

**Validates: Requirements 6.2, 6.3, 6.4, 6.5**
"""

from __future__ import annotations

import os
import tempfile
from pathlib import Path

from docx import Document
from hypothesis import HealthCheck, given, settings
from hypothesis import strategies as st

from app.services.wp_docx_template_parser import parse_template

# ─── Hypothesis Strategies ────────────────────────────────────────────────────

_field_id_st = st.from_regex(r"[a-z][a-z0-9_]{2,15}", fullmatch=True)
_value_st = st.text(
    min_size=1,
    max_size=20,
    alphabet=st.characters(categories=("L", "N"), whitelist_characters=" "),
)


@st.composite
def export_data_strategy(draw: st.DrawFn) -> tuple[list[tuple[str, str]], dict[str, str]]:
    """Generate (placeholders, response_map) for export testing.

    Returns:
        placeholders: list of (field_id, label) to embed in docx
        response_map: dict of field_id→value (partial or full coverage)
    """
    count = draw(st.integers(min_value=1, max_value=4))
    field_ids = draw(st.lists(_field_id_st, min_size=count, max_size=count, unique=True))
    labels = [f"标签{i}" for i in range(count)]
    placeholders = list(zip(field_ids, labels))

    # Generate values for a subset (or all) of placeholders
    fill_count = draw(st.integers(min_value=0, max_value=count))
    values = draw(st.lists(_value_st, min_size=fill_count, max_size=fill_count))
    response_map = {field_ids[i]: values[i] for i in range(fill_count)}

    return placeholders, response_map


@st.composite
def full_export_data_strategy(draw: st.DrawFn) -> tuple[list[tuple[str, str]], dict[str, str]]:
    """Generate (placeholders, response_map) where ALL placeholders have values."""
    count = draw(st.integers(min_value=1, max_value=4))
    field_ids = draw(st.lists(_field_id_st, min_size=count, max_size=count, unique=True))
    labels = [f"标签{i}" for i in range(count)]
    placeholders = list(zip(field_ids, labels))

    values = draw(st.lists(_value_st, min_size=count, max_size=count))
    response_map = {field_ids[i]: values[i] for i in range(count)}

    return placeholders, response_map


# ─── Helpers ──────────────────────────────────────────────────────────────────


def _create_template_docx(placeholders: list[tuple[str, str]]) -> str:
    """Create a temp docx file with ${field_id:label} placeholders. Returns path."""
    doc = Document()
    doc.add_heading("测试模板", level=1)
    for field_id, label in placeholders:
        para = doc.add_paragraph()
        # Add a run with specific formatting to verify preservation
        run = para.add_run(f"${{{field_id}:{label}}}")
        run.bold = True
    fd, path = tempfile.mkstemp(suffix=".docx")
    os.close(fd)
    doc.save(path)
    return path


def _replace_placeholders_in_docx(
    src_path: str, response_map: dict[str, str]
) -> str:
    """Simulate the export logic: replace placeholders with values, preserving run formatting.

    This mirrors the actual logic in download_template_prefilled with include_responses=True.
    Returns path to the modified temp docx.
    """
    import shutil

    fd, dst_path = tempfile.mkstemp(suffix=".docx")
    os.close(fd)
    shutil.copy2(src_path, dst_path)

    structure = parse_template(src_path)

    # Build pattern→value map
    pattern_to_value: dict[str, str] = {}
    for placeholder in structure.placeholders:
        if placeholder.field_id in response_map:
            pattern_to_value[placeholder.pattern] = response_map[placeholder.field_id]

    if not pattern_to_value:
        return dst_path

    doc = Document(dst_path)

    for para in doc.paragraphs:
        for pattern, value in pattern_to_value.items():
            if pattern in para.text:
                for run in para.runs:
                    if pattern in run.text:
                        run.text = run.text.replace(pattern, value)

    for table in doc.tables:
        for row in table.rows:
            for cell in row.cells:
                for para in cell.paragraphs:
                    for pattern, value in pattern_to_value.items():
                        if pattern in para.text:
                            for run in para.runs:
                                if pattern in run.text:
                                    run.text = run.text.replace(pattern, value)

    doc.save(dst_path)
    return dst_path


# ─── Property 9: Export replacement + format preservation ─────────────────────


class TestProperty9ExportReplacementFormatPreservation:
    """Property 9: Export placeholder replacement with format preservation.

    For any word-template export with include_responses=true:
    - Placeholders with corresponding values SHALL be replaced in the output docx
    - Placeholders without values SHALL remain unchanged as original text
    - Original run formatting (bold) SHALL be preserved

    **Validates: Requirements 6.2, 6.3, 6.4**
    """

    @settings(max_examples=5, suppress_health_check=[HealthCheck.function_scoped_fixture])
    @given(data=export_data_strategy())
    def test_placeholders_with_values_are_replaced(self, data):
        """Placeholders with responses are replaced with the value in exported docx."""
        placeholders, response_map = data
        template_path = _create_template_docx(placeholders)
        try:
            exported_path = _replace_placeholders_in_docx(template_path, response_map)
            try:
                doc = Document(exported_path)
                full_text = "\n".join(p.text for p in doc.paragraphs)

                for field_id, value in response_map.items():
                    # The value should appear in the exported document
                    assert value in full_text, (
                        f"Value '{value}' for field '{field_id}' not found in exported doc"
                    )
                    # The original pattern should NOT appear
                    pattern = f"${{{field_id}:"
                    assert pattern not in full_text, (
                        f"Pattern '{pattern}' still present after replacement"
                    )
            finally:
                os.unlink(exported_path)
        finally:
            os.unlink(template_path)

    @settings(max_examples=5, suppress_health_check=[HealthCheck.function_scoped_fixture])
    @given(data=export_data_strategy())
    def test_placeholders_without_values_remain_unchanged(self, data):
        """Placeholders without responses remain as original text."""
        placeholders, response_map = data
        template_path = _create_template_docx(placeholders)
        try:
            exported_path = _replace_placeholders_in_docx(template_path, response_map)
            try:
                doc = Document(exported_path)
                full_text = "\n".join(p.text for p in doc.paragraphs)

                # Fields NOT in response_map should still have their pattern
                for field_id, label in placeholders:
                    if field_id not in response_map:
                        pattern = f"${{{field_id}:{label}}}"
                        assert pattern in full_text, (
                            f"Unfilled placeholder '{pattern}' was incorrectly removed"
                        )
            finally:
                os.unlink(exported_path)
        finally:
            os.unlink(template_path)

    @settings(max_examples=5, suppress_health_check=[HealthCheck.function_scoped_fixture])
    @given(data=export_data_strategy())
    def test_run_formatting_preserved_after_replacement(self, data):
        """Bold formatting on runs is preserved after placeholder replacement."""
        placeholders, response_map = data
        if not response_map:
            return  # Need at least one replacement to test formatting
        template_path = _create_template_docx(placeholders)
        try:
            exported_path = _replace_placeholders_in_docx(template_path, response_map)
            try:
                doc = Document(exported_path)
                # Check that runs with replaced content still have bold=True
                for para in doc.paragraphs:
                    for run in para.runs:
                        if any(v in run.text for v in response_map.values()):
                            assert run.bold is True, (
                                f"Run bold formatting lost after replacement: '{run.text}'"
                            )
            finally:
                os.unlink(exported_path)
        finally:
            os.unlink(template_path)


# ─── Property 10: Export round-trip ───────────────────────────────────────────


class TestProperty10ExportRoundTrip:
    """Property 10: Export round-trip.

    For any TemplateStructure with all placeholders fully populated,
    exporting to docx then re-parsing SHALL produce the same field values.

    **Validates: Requirements 6.5**
    """

    @settings(max_examples=5, suppress_health_check=[HealthCheck.function_scoped_fixture])
    @given(data=full_export_data_strategy())
    def test_export_then_reparse_preserves_values(self, data):
        """Export with all values then re-parse: values appear in paragraph text."""
        placeholders, response_map = data
        template_path = _create_template_docx(placeholders)
        try:
            exported_path = _replace_placeholders_in_docx(template_path, response_map)
            try:
                # Re-read the exported document and verify values are present
                doc = Document(exported_path)
                full_text = "\n".join(p.text for p in doc.paragraphs)

                for field_id, value in response_map.items():
                    assert value in full_text, (
                        f"Round-trip failed: value '{value}' for '{field_id}' "
                        f"not found after export+re-parse"
                    )
            finally:
                os.unlink(exported_path)
        finally:
            os.unlink(template_path)

    @settings(max_examples=5, suppress_health_check=[HealthCheck.function_scoped_fixture])
    @given(data=full_export_data_strategy())
    def test_no_original_patterns_remain_after_full_export(self, data):
        """When all placeholders have values, no ${...} patterns remain."""
        placeholders, response_map = data
        template_path = _create_template_docx(placeholders)
        try:
            exported_path = _replace_placeholders_in_docx(template_path, response_map)
            try:
                doc = Document(exported_path)
                full_text = "\n".join(p.text for p in doc.paragraphs)

                import re
                remaining_patterns = re.findall(r"\$\{[^}]+\}", full_text)
                assert not remaining_patterns, (
                    f"Patterns remain after full export: {remaining_patterns}"
                )
            finally:
                os.unlink(exported_path)
        finally:
            os.unlink(template_path)

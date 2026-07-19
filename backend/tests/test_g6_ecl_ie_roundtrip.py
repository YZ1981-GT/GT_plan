"""G6 ECL 导入导出纯函数 round-trip（G6-14 双表 / G6-15 v2）。"""

from __future__ import annotations

import io

from openpyxl import load_workbook

from app.routers.wp_render_strategies._g6_other_bond_investment_ecl_import_export import (
    _G6_14_PRIMARY_ITEM_ID,
    _ITEM_IDS,
    _build_g6_14_multi_sheet_workbook,
    _build_g6_15_multi_sheet_workbook,
    _parse_g6_14_import,
    _parse_g6_15_import,
)


def test_g6_14_primary_and_compat_item_ids():
    assert _G6_14_PRIMARY_ITEM_ID == "G6-14-reversal-writeoff-data"
    assert _ITEM_IDS["G6-14"] == "G6-14-rows"


def test_g6_14_multi_sheet_export_import_round_trip():
    reversals = [{
        "seq": 1,
        "unitName": "农发债2024",
        "crossSheetInvestmentId": "inv-nfb-2024",
        "kind": "转回",
        "reversalReason": "信用风险改善",
        "recoveryMethod": "现金收回",
        "originalBasis": "原减值模型",
        "reversalAmount": 50000.0,
        "accumulatedProvision": 120000.0,
        "reasonAnalysis": "已核实现金流改善",
        "isReasonable": "合理",
        "indexRef": "G6-12",
    }]
    writeoffs = [{
        "seq": 1,
        "unitName": "企业债2023",
        "writeOffType": "公司债",
        "writeOffAmount": 200000.0,
        "writeOffReason": "债务人破产",
        "writeOffProcedure": "董事会决议",
        "isRelatedParty": True,
        "reasonAnalysis": "关联交易已专项复核",
        "isReasonable": "合理",
        "indexRef": "法务函",
    }]

    wb = _build_g6_14_multi_sheet_workbook(reversals, writeoffs)
    assert any("转回检查" in name for name in wb.sheetnames)
    assert any("核销检查" in name for name in wb.sheetnames)

    buf = io.BytesIO()
    wb.save(buf)
    content = buf.getvalue()

    parsed_reversals, parsed_writeoffs, errors = _parse_g6_14_import(content)
    assert errors == []
    assert len(parsed_reversals) == 1
    assert len(parsed_writeoffs) == 1
    assert parsed_reversals[0]["unitName"] == "农发债2024"
    assert parsed_reversals[0]["crossSheetInvestmentId"] == "inv-nfb-2024"
    assert parsed_reversals[0]["kind"] == "转回"
    assert parsed_reversals[0]["reversalAmount"] == 50000.0
    assert parsed_reversals[0]["accumulatedProvision"] == 120000.0
    assert parsed_writeoffs[0]["unitName"] == "企业债2023"
    assert parsed_writeoffs[0]["isRelatedParty"] is True
    assert parsed_writeoffs[0]["writeOffAmount"] == 200000.0


def test_g6_12_preserves_cross_sheet_investment_id():
    from app.routers.wp_render_strategies._g6_other_bond_investment_ecl_import_export import (
        _build_g6_12_multi_sheet_workbook,
        _parse_g6_12_import,
    )
    rows = [{
        "investProject": "国开债",
        "crossSheetInvestmentId": "inv-cdb-1",
        "amortizedCost": 100.0,
        "fairValue": 101.0,
        "creditLossRate": 0.01,
        "impairmentProvision": 1.0,
        "bookValue": 99.0,
        "balanceAdjustment": 0.0,
        "adjustedCreditLossRate": 0.01,
        "impairmentAdjustment": 0.0,
        "stage": "Stage1",
        "ociImpact": 0.0,
        "indexRef": "",
        "adjBalance": 100.0,
        "adjImpairment": 1.0,
        "adjBookValue": 99.0,
        "adjFairValue": 101.0,
        "priorImpairment": 0.0,
        "currentProvision": 1.0,
        "currentReversal": 0.0,
        "ociAdjustment": 0.0,
        "differenceNote": "",
    }]
    wb = _build_g6_12_multi_sheet_workbook(rows)
    buf = io.BytesIO()
    wb.save(buf)
    parsed, errors = _parse_g6_12_import(buf.getvalue())
    assert errors == []
    assert len(parsed) == 1
    assert parsed[0]["crossSheetInvestmentId"] == "inv-cdb-1"
    assert parsed[0]["id"] == "inv-cdb-1"


def test_g6_14_legacy_flat_sheet_still_imports():
    from openpyxl import Workbook

    wb = Workbook()
    ws = wb.active
    ws.title = "转回核销"
    ws.append(["G6-14 旧版单表"])
    ws.append([
        "序号", "投资项目", "转回/核销类型", "金额",
        "原因", "审批程序", "合理性结论", "索引",
    ])
    ws.append([1, "债A", "转回", 1000, "改善", "审批", "合理", "idx-1"])
    ws.append([2, "债B", "核销", 2000, "破产", "董事会", "合理", "idx-2"])
    buf = io.BytesIO()
    wb.save(buf)

    reversals, writeoffs, errors = _parse_g6_14_import(buf.getvalue())
    assert not any(e.get("reason", "").startswith("无法识别") for e in errors)
    assert len(reversals) == 1
    assert len(writeoffs) == 1
    assert reversals[0]["unitName"] == "债A"
    assert writeoffs[0]["unitName"] == "债B"


def test_g6_15_v2_criteria_and_period_round_trip():
    rows = [
        {
            "period": "occurrence",
            "date": "2024-12-15",
            "voucherNo": "PZ-001",
            "businessContent": "计提减值",
            "businessType": "公允价值变动",
            "counterAccount": "资产减值损失",
            "detailAccount": "1503",
            "debitAmount": 100000.0,
            "creditAmount": 0.0,
            "supportingDoc": "减值测算表",
            "checkOriginal": "Y",
            "checkAuthorized": "Y",
            "checkAccounting": "Y",
            "checkInitialCost": "NA",
            "checkInterest": "N",
            "checkFairValue": "Y",
            "indexRef": "G6-12",
            "manualAbnormal": False,
            "isAbnormal": True,
            "abnormalNote": "利息核对未通过",
            "riskLevel": "中",
            "suggestion": "补充测算底稿",
            "remark": "已沟通",
            "source": "抽样",
            "selectionReason": "大额",
            "samplingMethod": "MUS",
            "selectionCategory": "specific",
            "sourceId": "s-1",
            "attachment": "att.pdf",
            "attachmentId": "a-1",
        },
        {
            "period": "post",
            "date": "2025-01-10",
            "voucherNo": "PZ-100",
            "businessContent": "期后处置",
            "businessType": "处置",
            "counterAccount": "银行存款",
            "detailAccount": "",
            "debitAmount": 0.0,
            "creditAmount": 50000.0,
            "supportingDoc": "交割单",
            "checkOriginal": "Y",
            "checkAuthorized": "Y",
            "checkAccounting": "Y",
            "checkInitialCost": "Y",
            "checkInterest": "NA",
            "checkFairValue": "Y",
            "indexRef": "",
            "manualAbnormal": False,
            "isAbnormal": False,
            "abnormalNote": "",
            "riskLevel": "低",
            "suggestion": "",
            "remark": "",
            "source": "手工",
            "selectionReason": "",
            "samplingMethod": "",
            "selectionCategory": "manual",
            "sourceId": "",
            "attachment": "",
            "attachmentId": "",
        },
    ]
    criteria = {
        "populationDebitCount": 12,
        "populationDebitAmount": 1000000,
        "samplingMethod": "MUS",
        "sampleSize": 5,
        "bookDebitOccurrence": 800000,
        "bookCreditOccurrence": 200000,
    }

    wb = _build_g6_15_multi_sheet_workbook(rows, criteria=criteria)
    assert "样本标准" in wb.sheetnames
    assert "本期发生额" in wb.sheetnames
    assert "期后处置新增" in wb.sheetnames

    buf = io.BytesIO()
    wb.save(buf)
    parsed_rows, errors, parsed_criteria = _parse_g6_15_import(buf.getvalue())

    assert errors == [] or all("截断" not in e.get("reason", "") for e in errors)
    assert parsed_criteria.get("samplingMethod") == "MUS"
    assert parsed_criteria.get("sampleSize") == 5
    assert parsed_criteria.get("populationDebitCount") == 12

    by_voucher = {r["voucherNo"]: r for r in parsed_rows}
    assert set(by_voucher) >= {"PZ-001", "PZ-100"}
    occ = by_voucher["PZ-001"]
    assert occ["period"] == "occurrence"
    assert occ["checkInterest"] == "N"
    assert occ["checkInitialCost"] == "NA"
    assert occ["isAbnormal"] is True
    assert occ["abnormalNote"] == "利息核对未通过"
    assert occ["selectionCategory"] == "specific"
    assert occ["attachmentId"] == "a-1"

    post = by_voucher["PZ-100"]
    assert post["period"] == "post"
    assert post["creditAmount"] == 50000.0


def test_g6_14_template_has_dual_tabs():
    wb = _build_g6_14_multi_sheet_workbook([], [], template_only=True)
    names = wb.sheetnames
    assert any("转回检查" in n for n in names)
    assert any("核销检查" in n for n in names)
    ws = next(wb[n] for n in names if "转回检查" in n)
    headers = [c.value for c in next(ws.iter_rows(min_row=2, max_row=2))]
    assert "类型(转回/收回)" in headers
    assert "收回或转回金额" in headers
    assert "跨表投资ID" in headers

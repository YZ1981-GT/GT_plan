"""附注种子 JSON ↔ 四版 Word：章节树 / 表格比对工具自检 + 真实文件守卫。

spec: disclosure-note-template-word-realignment Task 1 / 2（design §2）

- 自检用例用合成 docx + 合成种子，证明每类差异都能检出（防止比对器空转）；
- 真实文件守卫在 Phase 2 修复前标 ``xfail(strict=True)``：差异真实存在，不把现状写进白名单冒充通过。
"""

from __future__ import annotations

from pathlib import Path

import docx
import pytest

from app.services.note_word_tree import (
    ALLOWLIST_PATH,
    build_report,
    check_against_allowlist,
    compare_variant,
    load_seed,
    match_variant,
    parse_word_docx,
)


def _docx(path: Path, blocks: list[tuple]) -> Path:
    """blocks: ("h", 层级, 文字) / ("p", 文字) / ("t", 表头行列表, 数据行列表[, 合并说明])。"""
    document = docx.Document()
    for block in blocks:
        if block[0] == "h":
            document.add_heading(block[2], level=block[1])
        elif block[0] == "p":
            document.add_paragraph(block[1])
        else:
            header_rows, data_rows = block[1], block[2]
            rows = header_rows + data_rows
            table = document.add_table(rows=len(rows), cols=len(rows[0]))
            for r, row in enumerate(rows):
                for c, text in enumerate(row):
                    table.cell(r, c).text = text
            for (r0, c0, r1, c1, text) in (block[3] if len(block) > 3 else []):
                merged = table.cell(r0, c0).merge(table.cell(r1, c1))
                merged.text = text
    document.save(str(path))
    return path


def _sec(code, title, parent=None, *, level=2, order=0, scope="both", tables=None, sid=None):
    return {
        "section_number": code, "section_title": title, "level": level, "scope": scope,
        "parent_section_id": parent, "section_id": sid or f"id-{code}", "sort_order": order,
        "tables": tables or [],
    }


SEED = {
    "sections": [
        _sec("一", "基本情况", level=1, order=0, sid="c1"),
        _sec("一、1", "基本情况", "c1", order=1),
        _sec("一、所得税费用", "所得税费用", "c1", order=2),  # 放错章：Word 在项目注释章
        _sec("二", "项目注释", level=1, order=10, sid="c2"),
        _sec("二、1", "货币资金", "c2", order=11, tables=[
            {"name": "货币资金", "headers": ["项目", "期末余额"], "rows": [{"label": "库存现金"}, {"label": "合计"}]},
        ]),
        _sec("二、2", "【企业应当披露下列信息：", "c2", order=12),  # 提示语误做成的节
        _sec("二、3", "应收账款", "c2", order=13, scope="consolidated_only"),
        _sec("二、4", "营业收入与营业成本", "c2", order=14, tables=[
            {"name": "营业收入", "headers": ["项目", "本期收入"], "rows": [{"label": "主营业务"}]},
        ]),
        _sec("二、5", "其他收益", "c2", order=15, tables=[{"name": "其他收益", "headers": ["项目", "本期发生额"]}]),
        _sec("三", "承诺事项", level=1, order=20, sid="c3"),
        _sec("三、1", "财务报表之批准", "c3", order=21),  # 章名错、章正文叶子用 Word 章名
    ]
}


def _listed_word(tmp_path: Path) -> Path:
    return _docx(tmp_path / "listed_standalone.docx", [
        ("h", 1, "基本情况"), ("p", "本公司是一家股份有限公司。"), ("p", "{{seq:一}}"),
        ("h", 1, "项目注释"),
        ("h", 2, "应收账款"),
        ("h", 2, "货币资金"), ("t", [["项目", "期末余额", "期初余额"]], [["库存现金", "1", "2"], ["合计", "1", "2"]]),
        ("h", 2, "营业收入和营业成本（披露格式参考附注五、62）"),
        ("t", [["项目", "本期发生额"]], [["主营业务", ""], ["其他业务", ""]]),
        ("h", 2, "所得税费用"),
        ("h", 2, "预付款项"),
        ("h", 2, "其他收益"),
        ("h", 1, "财务报表之批准"), ("p", "本财务报表业经董事会批准。"),
    ])


def _keys(diffs) -> set[tuple[str, str]]:
    return {(d.kind, d.key) for d in diffs}


def test_every_diff_kind_is_detected(tmp_path):
    word = parse_word_docx(_listed_word(tmp_path), "listed", variant="listed_standalone")
    diffs = compare_variant(match_variant(word, load_seed(SEED), "listed_standalone"))
    assert _keys(diffs) == {
        ("chapter_title", "三"),  # 章名错，靠章正文叶子配对
        ("misplaced", "一、所得税费用"),
        ("json_only", "二、2"),
        ("scope", "二、3"),  # 单体 Word 有、种子标 consolidated_only
        ("section_title", "二、4"),  # 同章同空档，措辞不同（与 / 和）
        ("order", "二、1"),
        ("word_only", "项目注释/预付款项"),
        ("table_columns", "二、1#t0"),
        ("table_count", "二、5"),  # 种子 1 张、Word 0 张
        ("table_headers", "二、4#t0"),  # 列数相同、列名不同
        ("table_rows", "二、4#t0"),  # Word 多一行「其他业务」
    }, sorted(_keys(diffs))
    detail = {d.key: d.detail for d in diffs}
    assert detail["二、2"].startswith("guidance")
    assert "第 2 章「项目注释」" in detail["一、所得税费用"]


def test_chapter_body_leaves_and_markers_are_not_diffs(tmp_path):
    word = parse_word_docx(_listed_word(tmp_path), "listed", variant="listed_standalone")
    match = match_variant(word, load_seed(SEED), "listed_standalone")
    leaves = {match.sections_by_position()[pos].code for pos in match.body_leaves}
    assert leaves == {"一、1", "三、1"}
    assert word.chapters[0].body == ["本公司是一家股份有限公司。"]  # 模板标记 {{seq:一}} 不算正文


def test_soe_section_level_is_decided_per_chapter_and_merged_headers(tmp_path):
    path = _docx(tmp_path / "soe.docx", [
        ("h", 1, "企业合并及合并财务报表"),
        ("h", 2, "重要非全资子公司情况"),
        ("h", 3, "少数股东"), ("t", [["企业名称", "持股比例"]], [["", ""]]),
        ("h", 3, "主要财务信息"), ("t", [["项目", "期末数"]], [["流动资产", ""]]),
        ("h", 2, "本期新纳入合并范围的主体"),
        ("h", 1, "财务报表主要项目注释"),
        ("p", "##SECTION:八、1##"),
        ("h", 3, "X、货币资金"),
        ("h", 4, "（1）明细"),
        ("t", [["项目", "", ""], ["", "账面余额", "减值准备"]], [["应收A", "1", "2"], ["……", "", ""]],
         [(0, 0, 1, 0, "项目"), (0, 1, 0, 2, "期末余额")]),
        # 首列表头未纵向合并、下格留白（真实模板「账龄」列即如此）：留白不算一级表头
        ("t", [["账龄", "", ""], ["", "金额", "比例"]], [["1年以内", "1", "2"]], [(0, 1, 0, 2, "期末余额")]),
    ])
    word = parse_word_docx(path, "soe")
    merge, notes = word.chapters
    assert [s.title for s in merge.sections] == ["重要非全资子公司情况", "本期新纳入合并范围的主体"]
    assert merge.sections[0].subtitles == ["少数股东", "主要财务信息"]
    assert len(merge.sections[0].tables) == 2  # H3 小标题下的表归所属 H2 节
    (section,) = notes.sections
    assert section.norm == "货币资金"  # 模板占位序号「X、」不进比对键
    table = section.tables[0]
    assert (table.columns, table.header_rows) == (3, 2)
    assert table.leaf_headers == ["项目", "账面余额", "减值准备"]
    assert table.groups == [None, "期末余额", "期末余额"]
    assert table.row_labels == ["应收A", "……"]
    aging = section.tables[1]
    assert aging.leaf_headers == ["账龄", "金额", "比例"] and aging.groups == [None, "期末余额", "期末余额"]


def test_allowlist_reports_unexpected_stale_and_requires_reason(tmp_path):
    word = parse_word_docx(_listed_word(tmp_path), "listed", variant="listed_standalone")
    diffs = compare_variant(match_variant(word, load_seed(SEED), "listed_standalone"))
    allow = [
        {"variant": "listed_standalone", "kind": "order", "key": "二、1", "reason": "测试"},
        {"variant": "listed_standalone", "kind": "order", "key": "不存在", "reason": "测试"},
    ]
    unexpected, stale = check_against_allowlist(diffs, allow)
    assert ("order", "二、1") not in _keys(unexpected) and len(unexpected) == len(diffs) - 1
    assert [e["key"] for e in stale] == ["不存在"]
    with pytest.raises(ValueError, match="缺理由"):
        check_against_allowlist(diffs, [{"variant": "v", "kind": "k", "key": "x", "reason": " "}])


def _real_check(kind_filter):
    import json

    allowlist = json.loads(ALLOWLIST_PATH.read_text(encoding="utf-8")) if ALLOWLIST_PATH.exists() else []
    diffs = [d for d in build_report() if kind_filter(d.kind)]
    allowlist = [e for e in allowlist if kind_filter(e["kind"])]
    return check_against_allowlist(diffs, allowlist)


def test_seed_tree_matches_four_word_templates():
    """章节树（章 / 节 / 归属 / 顺序 / scope / 标题）与四版 Word 一致；例外只在白名单（理由必填）。"""
    unexpected, stale = _real_check(lambda kind: not kind.startswith("table_"))
    assert not stale, stale
    assert not unexpected, f"{len(unexpected)} 条白名单外差异，例：{unexpected[:5]}"


@pytest.mark.xfail(strict=True, reason="表格以 Word 还是底稿模板为准待用户拍板（tasks.md「比对基线」结论第 3 条）")
def test_seed_tables_match_four_word_templates():
    unexpected, stale = _real_check(lambda kind: kind.startswith("table_"))
    assert not stale, stale
    assert not unexpected, f"{len(unexpected)} 条白名单外差异，例：{unexpected[:5]}"


def test_logical_columns_merge_grid_columns_split_only_in_body():
    from app.services.note_word_table import _logical_columns

    # 表头 2 行：第 0 行「期末」跨网格列 1~3、第 1 行「余额」跨 2~3 ⇒ 网格列 2、3 在表头里始终同格
    grid = {
        (0, 0): (0, 0), (0, 1): (0, 1), (0, 2): (0, 1), (0, 3): (0, 1),
        (1, 0): (0, 0), (1, 1): (1, 1), (1, 2): (1, 2), (1, 3): (1, 2),
        (2, 0): (2, 0), (2, 1): (2, 1), (2, 2): (2, 2), (2, 3): (2, 3),  # 数据行把 2、3 拆开
    }
    assert _logical_columns(4, 2, grid) == [0, 1, 2]


def test_placeholder_rows_are_ignored_in_row_comparison():
    from app.services.note_word_table import _row_keys

    assert _row_keys(["库存现金", "……", "......", "可无限量添加行", "1", "10", "合 计：", None, ""]) == ["库存现金", "合计"]

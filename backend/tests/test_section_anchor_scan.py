"""锚点区间扫描与块定位降级链测试 — deliverable-lineage-wiring-and-writeback-closure Task 1.1

覆盖：
- Property 4：锚点写入 ↔ 扫描互逆（标记清理后仍可定位）
- Property 5：块定位降级链单调（anchor / marker / none 三态均不抛异常）
- Rendered_Block_Hash 口径（排除 ## 标记行、规范化、对改动敏感）
- 反向自检：内容控件开启后仅看 body 级会把整节看成单个 sdt（故必须展平）

本层用合成 docx 是合法的（design §Testing Strategy 第 1 层）；**真实链路验收在 Task 11**。
"""

from __future__ import annotations

from io import BytesIO

import pytest
from docx import Document
from docx.oxml import OxmlElement
from docx.oxml.ns import qn
from hypothesis import given, settings
from hypothesis import strategies as st

from app.services.section_anchor_utils import (
    SectionBlock,
    anchor_block_content_host,
    anchor_name,
    block_text_hash,
    block_text_of,
    normalize_block_text,
    resolve_section_blocks,
    scan_anchor_blocks,
    write_section_anchors,
)
from app.services.word_doc_utils import remove_section_markers, scan_section_blocks

ALL_CODES = ["八、1", "八、2", "五、1", "八、22"]


# ─── 夹具 ────────────────────────────────────────────────────────────────────


def _build_doc(sections: dict[str, list[str]]) -> Document:
    """构建含 ``##SECTION:code##`` 块的合成 docx（每节若干正文段落）。"""
    doc = Document()
    for code, lines in sections.items():
        doc.add_paragraph(f"##SECTION:{code}##")
        for line in lines:
            doc.add_paragraph(line)
        doc.add_paragraph(f"##/SECTION:{code}##")
    return doc


def _roundtrip(doc: Document) -> Document:
    """序列化 → 反序列化（模拟落盘 / OnlyOffice 往返后再由 python-docx 读回）。"""
    buf = BytesIO()
    doc.save(buf)
    buf.seek(0)
    return Document(buf)


def _write_anchors(doc: Document, kept_codes: list[str]) -> dict[str, str]:
    blocks = {b.section_code: b for b in scan_section_blocks(doc)}
    kept_blocks = [
        SectionBlock(
            section_code=code,
            open_el=blocks[code].open_el,
            close_el=blocks[code].close_el,
        )
        for code in kept_codes
        if code in blocks
    ]
    return write_section_anchors(doc, kept_blocks)


# WordprocessingML 主命名空间前缀（python-docx 元素标签）。
W = "{http://schemas.openxmlformats.org/wordprocessingml/2006/main}"


def _oo_saved(doc):
    """模拟 OnlyOffice 在线编辑保存：body 级 ``sec_`` 书签被挪进块首段落开头 / 块尾段落末尾
    （安徽附注交付件 v15 实测 106/106 节；内容控件开时落在控件里的首 / 末段落）。

    夹具辅助：原在 tests.services.test_deliverable_block_transplant，该文件随 block_transplant
    架构弃用而移除，这里内联保留（只构造 docx 形态，不依赖任何业务逻辑）。
    """
    body = doc.element.body
    containers = (f"{W}p", f"{W}tbl", f"{W}sdt")
    starts = [e for e in body.iterchildren(f"{W}bookmarkStart") if (e.get(qn("w:name")) or "").startswith("sec_")]
    ids = {e.get(qn("w:id")) for e in starts}
    for start in starts:
        nxt = start.getnext()
        while nxt.tag not in containers:
            nxt = nxt.getnext()
        first_p = nxt if nxt.tag == f"{W}p" else next(nxt.iter(f"{W}p"))
        body.remove(start)
        ppr = first_p.find(f"{W}pPr")
        if ppr is not None:
            ppr.addnext(start)
        else:
            first_p.insert(0, start)
    for end in [e for e in body.iterchildren(f"{W}bookmarkEnd") if e.get(qn("w:id")) in ids]:
        prev = end.getprevious()
        while prev.tag not in containers:
            prev = prev.getprevious()
        last_p = prev if prev.tag == f"{W}p" else list(prev.iter(f"{W}p"))[-1]
        body.remove(end)
        last_p.append(end)
    return doc


def _all_text(doc: Document) -> list[str]:
    return [p.text for p in doc.paragraphs]


# ─── Property 4：写入 ↔ 扫描互逆 ─────────────────────────────────────────────


@given(kept=st.lists(st.sampled_from(ALL_CODES), min_size=1, max_size=4, unique=True))
@settings(max_examples=5, deadline=None)
def test_property_4_anchor_scan_is_inverse_of_write(kept):
    # Feature: deliverable-lineage-wiring-and-writeback-closure, Property 4
    """写锚点 → 清标记 → 往返 → scan_anchor_blocks 得到的 code 集合 == kept。"""
    sections = {code: [f"{code} 正文一", f"{code} 正文二"] for code in ALL_CODES}
    doc = _build_doc(sections)
    _write_anchors(doc, kept)
    remove_section_markers(doc)
    doc = _roundtrip(doc)

    blocks = scan_anchor_blocks(doc)
    assert {b.section_code for b in blocks} == set(kept)
    for b in blocks:
        assert b.anchor_name == anchor_name(b.section_code)
        # 区间内容元素非空且不含标记行
        texts = block_text_of(b.elements)
        assert b.section_code in texts
        assert "##" not in texts


def test_anchor_scan_covers_exactly_its_own_content():
    """锚点区间只覆盖本章节内容，不串到邻节。"""
    doc = _build_doc({"八、1": ["甲"], "八、2": ["乙"]})
    _write_anchors(doc, ["八、1", "八、2"])
    remove_section_markers(doc)
    doc = _roundtrip(doc)

    by_code = {b.section_code: b for b in scan_anchor_blocks(doc)}
    assert block_text_of(by_code["八、1"].elements) == "甲"
    assert block_text_of(by_code["八、2"].elements) == "乙"


def test_trimmed_sections_have_no_anchor():
    """裁剪章节不写锚点 → 扫描结果里没有它（Property 1 的扫描侧印证）。"""
    doc = _build_doc({"八、1": ["甲"], "八、2": ["乙"]})
    _write_anchors(doc, ["八、1"])
    remove_section_markers(doc)
    doc = _roundtrip(doc)

    assert {b.section_code for b in scan_anchor_blocks(doc)} == {"八、1"}


# ─── 内容控件并存 ────────────────────────────────────────────────────────────


def test_scan_works_with_content_controls_and_flattens():
    """内容控件开启后：body 级只剩一个 w:sdt，扫描须下钻取到内部段落。"""
    from app.services.content_control_injector import inject_content_controls_for_blocks

    doc = _build_doc({"八、1": ["甲一", "甲二"]})
    _write_anchors(doc, ["八、1"])
    inject_content_controls_for_blocks(doc, scan_section_blocks(doc))
    remove_section_markers(doc)
    doc = _roundtrip(doc)

    blocks = scan_anchor_blocks(doc)
    assert len(blocks) == 1
    b = blocks[0]
    # 反向自检：body 级容器是单个 sdt（只看 containers 取不到文字）
    assert len(b.containers) == 1
    assert b.containers[0].tag == qn("w:sdt")
    assert block_text_of(b.containers) == "", "containers 未展平时不应有文字（证明展平必要）"
    # 展平后拿得到内部两段
    assert block_text_of(b.elements) == "甲一\n甲二"


def test_content_host_prefers_sdt_content():
    """有内容控件时插入宿主是 sdtContent；无控件时是 body。"""
    from app.services.content_control_injector import inject_content_controls_for_blocks

    doc = _build_doc({"八、1": ["甲"]})
    _write_anchors(doc, ["八、1"])
    remove_section_markers(doc)
    doc = _roundtrip(doc)
    block = scan_anchor_blocks(doc)[0]
    host, children = anchor_block_content_host(block)
    assert host.tag == qn("w:body")
    assert len(children) == 1

    doc2 = _build_doc({"八、1": ["甲"]})
    _write_anchors(doc2, ["八、1"])
    inject_content_controls_for_blocks(doc2, scan_section_blocks(doc2))
    remove_section_markers(doc2)
    doc2 = _roundtrip(doc2)
    block2 = scan_anchor_blocks(doc2)[0]
    host2, children2 = anchor_block_content_host(block2)
    assert host2.tag == qn("w:sdtContent")
    assert len(children2) == 1


# ─── 健壮性 ──────────────────────────────────────────────────────────────────


def test_non_section_bookmarks_ignored():
    """非 sec_ 前缀书签（Word 内部 / note bookmark）不得被当章节锚点。"""
    doc = _build_doc({"八、1": ["甲"]})
    body = doc.element.body
    bm = OxmlElement("w:bookmarkStart")
    bm.set(qn("w:id"), "9001")
    bm.set(qn("w:name"), "note_八_1")
    body.insert(0, bm)
    bm_end = OxmlElement("w:bookmarkEnd")
    bm_end.set(qn("w:id"), "9001")
    body.append(bm_end)

    assert scan_anchor_blocks(doc) == []


def test_unclosed_anchor_skipped():
    """bookmarkStart 无匹配 end → 跳过，不抛异常。"""
    doc = _build_doc({"八、1": ["甲"]})
    body = doc.element.body
    bm = OxmlElement("w:bookmarkStart")
    bm.set(qn("w:id"), "1000")
    bm.set(qn("w:name"), anchor_name("八、1"))
    body.insert(0, bm)

    assert scan_anchor_blocks(doc) == []


def test_duplicate_section_code_takes_first_only():
    """同 section_code 多锚点 → 取首个，禁止静默合并成一节。"""
    doc = _build_doc({"八、1": ["甲"]})
    _write_anchors(doc, ["八、1"])
    # 再手工追加一段同名锚点区间
    body = doc.element.body
    bm = OxmlElement("w:bookmarkStart")
    bm.set(qn("w:id"), "2000")
    bm.set(qn("w:name"), anchor_name("八、1"))
    body.append(bm)
    p = doc.add_paragraph("重复节内容")
    bm_end = OxmlElement("w:bookmarkEnd")
    bm_end.set(qn("w:id"), "2000")
    body.append(bm_end)
    remove_section_markers(doc)

    blocks = scan_anchor_blocks(doc)
    assert len(blocks) == 1
    assert "重复节内容" not in block_text_of(blocks[0].elements)
    assert p is not None


def test_scan_on_empty_document_returns_empty():
    assert scan_anchor_blocks(Document()) == []


# ─── Property 5：降级链单调 ──────────────────────────────────────────────────


def test_property_5_locate_mode_anchor():
    # Feature: deliverable-lineage-wiring-and-writeback-closure, Property 5
    doc = _build_doc({"八、1": ["甲"]})
    _write_anchors(doc, ["八、1"])
    remove_section_markers(doc)
    mode, blocks = resolve_section_blocks(_roundtrip(doc))
    assert mode == "anchor"
    assert [b.section_code for b in blocks] == ["八、1"]


def test_property_5_locate_mode_marker():
    # Feature: deliverable-lineage-wiring-and-writeback-closure, Property 5
    """存量 / 中间态文档：标记仍在、无锚点 → 回退 marker。"""
    doc = _build_doc({"八、1": ["甲"]})
    mode, blocks = resolve_section_blocks(doc)
    assert mode == "marker"
    assert [b.section_code for b in blocks] == ["八、1"]


def test_property_5_locate_mode_none():
    # Feature: deliverable-lineage-wiring-and-writeback-closure, Property 5
    """接线前生成的存量交付件：标记已清、无锚点 → none 且空列表（不抛）。"""
    doc = _build_doc({"八、1": ["甲"]})
    remove_section_markers(doc)
    mode, blocks = resolve_section_blocks(_roundtrip(doc))
    assert mode == "none"
    assert blocks == []


def test_locate_mode_anchor_wins_over_marker():
    """锚点与标记同时存在（中间态）→ 优先 anchor。"""
    doc = _build_doc({"八、1": ["甲"]})
    _write_anchors(doc, ["八、1"])
    mode, _ = resolve_section_blocks(doc)
    assert mode == "anchor"


@pytest.mark.parametrize(
    "builder",
    [
        lambda: Document(),
        lambda: _build_doc({}),
        lambda: _build_doc({"八、1": []}),
    ],
)
def test_locate_never_raises_on_degenerate_docs(builder):
    mode, blocks = resolve_section_blocks(builder())
    assert mode in ("anchor", "marker", "none")
    assert isinstance(blocks, list)


# ─── Rendered_Block_Hash 口径 ────────────────────────────────────────────────


def test_block_text_excludes_markers_and_tables():
    doc = _build_doc({"八、1": ["甲"]})
    block = scan_section_blocks(doc)[0]
    # 块内含开闭标记段落，提取结果不得包含
    assert block_text_of(block.elements) == "甲"


def test_normalize_block_text_collapses_whitespace_and_blank_lines():
    assert normalize_block_text("  甲   乙 \r\n\r\n  丙  ") == "甲 乙\n丙"
    assert normalize_block_text("") == ""


def test_block_text_hash_is_deterministic_and_sensitive():
    doc_a = _build_doc({"八、1": ["甲", "乙"]})
    doc_b = _build_doc({"八、1": ["甲 ", " 乙"]})
    doc_c = _build_doc({"八、1": ["甲", "丙"]})
    ha = block_text_hash(scan_section_blocks(doc_a)[0].elements)
    hb = block_text_hash(scan_section_blocks(doc_b)[0].elements)
    hc = block_text_hash(scan_section_blocks(doc_c)[0].elements)
    assert ha == hb, "仅空白差异不应改变哈希"
    assert ha != hc, "文字改动必须改变哈希"
    assert len(ha) == 64


def test_block_text_hash_differs_from_snapshot_hash_domain():
    """反向自检：Rendered_Block_Hash 与 source_snapshot_hash 是两个哈希域。

    历史缺陷正是拿后者去比块内文字 ⇒ 永远不等。本断言钉住「二者不可互换」，
    防后来者又把 `_detect_user_edits` 改回去比 source_snapshot_hash。
    """
    from app.services.deliverable_section_state_service import (
        compute_snapshot_hash_from_parts,
    )

    doc = _build_doc({"八、1": ["甲"]})
    block_hash = block_text_hash(scan_section_blocks(doc)[0].elements)
    snapshot_hash = compute_snapshot_hash_from_parts(
        section_code="八、1",
        text_content="甲",
        table_data=None,
        audited_amounts=[],
    )
    assert block_hash != snapshot_hash


# ─── 已知章节码精确反解（spec deliverable-trio-generation-accuracy Task 16b 追加） ───────────


MULTI = "九、1、或有负债"  # 第二个「、」也写成「_」⇒ 逆映射有损（真实交付件：安徽 2 节、测试集团甲 3 节）


def _anchored(codes: list[str]) -> Document:
    doc = _build_doc({c: [f"{c} 正文"] for c in codes})
    _write_anchors(doc, codes)
    remove_section_markers(doc)
    return _roundtrip(doc)


def test_lossy_reverse_mapping_is_real():
    """前提：不给已知章节码时，多「、」章节反解出的不是真实章节码（否则下面的判据空转）。"""
    codes = [b.section_code for b in scan_anchor_blocks(_anchored([MULTI, "八、1"]))]
    assert MULTI not in codes and "九、1_或有负债" in codes and "八、1" in codes


def test_known_codes_restore_exact_section_code():
    doc = _anchored([MULTI, "八、1", "五、12·1"])
    blocks = {b.section_code: b for b in scan_anchor_blocks(doc, known_codes=[MULTI, "八、1", "五、12·1"])}
    assert set(blocks) == {MULTI, "八、1", "五、12·1"}
    assert block_text_of(blocks[MULTI].elements) == f"{MULTI} 正文"
    # 统一入口同口径；已知码里没有的锚点照旧走逆映射，已知码不在文档里的不凭空造块
    mode, resolved = resolve_section_blocks(doc, known_codes=[MULTI, "八、99"])
    assert mode == "anchor"
    assert sorted(b.section_code for b in resolved) == sorted([MULTI, "八、1", "五、12_1"])


@pytest.mark.parametrize("controls", [True, False], ids=["内容控件开", "内容控件关"])
def test_bookmarks_moved_into_paragraphs_by_online_edit_are_scanned(controls):
    """OnlyOffice 在线编辑保存后 body 级书签被挪进块首 / 块尾段落（安徽附注 v15 实测 106/106 节）：
    原实现只认 body 级书签 ⇒ 0 个锚点，刷新 / 回填对在线编辑过的交付件全部失效。"""
    from app.services.content_control_injector import inject_content_controls_for_blocks

    sections = {"八、1": ["甲一", "甲二"], MULTI: ["或有负债正文"], "八、2": ["乙"]}
    doc = _build_doc(sections)
    _write_anchors(doc, list(sections))
    if controls:
        inject_content_controls_for_blocks(doc, scan_section_blocks(doc))
    remove_section_markers(doc)
    flat_blocks = scan_anchor_blocks(doc, known_codes=list(sections))
    expected = {b.section_code: block_text_of(b.elements) for b in flat_blocks}
    flat = {b.section_code: block_text_hash(b.elements) for b in flat_blocks}
    assert not any(b.nested for b in flat_blocks)

    saved = _roundtrip(_oo_saved(doc))
    assert not [e for e in saved.element.body.iterchildren(qn("w:bookmarkStart"))], "夹具：书签应已不在 body 级"
    mode, blocks = resolve_section_blocks(saved, known_codes=list(sections))
    assert mode == "anchor"
    assert {b.section_code: block_text_of(b.elements) for b in blocks} == expected
    assert all(b.nested for b in blocks)
    # 块哈希与 body 级形态一致：在线编辑保存本身不被当成人工修改
    assert {b.section_code: block_text_hash(b.elements) for b in blocks} == flat

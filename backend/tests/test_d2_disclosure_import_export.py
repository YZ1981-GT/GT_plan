"""D2 附注披露导入导出（sheet=D2-disc-listed/soe）纯函数测试

覆盖：
  P1 导出→导入往返：简单行表 / 组合计提父子结构 / 说明文本逐字一致
  P2 变体门控：国企专有表（其他组合方法 / 终止确认）只在 soe 出现
  P3 空表不清空：空白模板导入不产生 storage 写入（只给 warning）
  P4 账龄未匹配：组合分表账龄 label 不在当前段 → 跳过并 warning
  P5 说明子节名不识别 → 跳过并 warning；空说明不覆盖
  P6 sheet 参数解析：非法披露 sheet → 400
"""
import io

import pytest
from fastapi import HTTPException
from openpyxl import load_workbook

from app.routers.wp_render_strategies._d2_disclosure_import_export import (
    DISCLOSURE_SHEETS,
    WS_DERECOGNIZED,
    WS_NOTES,
    WS_OTHER_PORTFOLIO,
    WS_PORTFOLIO,
    build_disclosure_workbook,
    parse_disclosure_workbook,
    resolve_variant,
)

SEGMENTS = [
    {"key": "within1", "label": "1年以内", "dayFrom": 0},
    {"key": "y1to2", "label": "1至2年", "dayFrom": 366},
    {"key": "over3", "label": "3年以上", "dayFrom": 1096},
]


def _reload(wb):
    buf = io.BytesIO()
    wb.save(buf)
    buf.seek(0)
    return load_workbook(buf, data_only=True)


def _sample_data():
    return {
        "individual-rows": [
            {
                "name": "甲公司",
                "endAmount": 100.0,
                "priorAmount": 90.0,
                "provision": 10.0,
                "aging": "1至2年",
                "lossRate": 10.0,
                "basis": "已进入诉讼",
            }
        ],
        "portfolios": [
            {
                "groupId": "g1",
                "name": "应收中央企业客户",
                "rows": [
                    {"key": "within1", "label": "1年以内", "endAmount": 500.0, "priorAmount": 400.0},
                    {"key": "y1to2", "label": "1至2年", "endAmount": 50.0, "priorAmount": 40.0},
                ],
            }
        ],
        "reversal-rows": [
            {
                "companyName": "乙公司",
                "reversalReason": "债务人恢复偿付能力",
                "recoveryMethod": "银行转账",
                "originalBasis": "账龄组合",
                "cumulativeProvision": 20.0,
                "amount": 30.0,
            }
        ],
        "writeoff-rows": [
            {
                "companyName": "丙公司",
                "nature": "货款",
                "amount": 15.0,
                "reason": "破产清算",
                "procedure": "已经管理层审批",
                "relatedParty": "否",
            }
        ],
        "top5-rows": [
            {"companyName": "丁公司", "arAmount": 1000.0, "contractAssetAmount": 200.0, "provision": 30.0}
        ],
        "other-portfolio-rows": [{"name": "余额百分比组合", "endAmount": 60.0, "priorAmount": 50.0}],
        "derecognized-rows": [{"companyName": "戊公司", "amount": 70.0, "gainLoss": -5.0}],
        "notes": {"aging": "账龄结构说明", "top5": "前五名说明"},
    }


class TestVariantGating:
    def test_soe_only_sheets_present_only_in_soe(self):
        listed = build_disclosure_workbook("listed", SEGMENTS, None)
        soe = build_disclosure_workbook("soe", SEGMENTS, None)
        assert WS_OTHER_PORTFOLIO not in listed.sheetnames
        assert WS_DERECOGNIZED not in listed.sheetnames
        assert WS_OTHER_PORTFOLIO in soe.sheetnames
        assert WS_DERECOGNIZED in soe.sheetnames
        # 两个变体都有组合分表与说明文本
        for wb in (listed, soe):
            assert WS_PORTFOLIO in wb.sheetnames
            assert WS_NOTES in wb.sheetnames

    def test_resolve_variant(self):
        assert resolve_variant("D2-disc-listed") == "listed"
        assert resolve_variant("D2-disc-soe") == "soe"
        assert set(DISCLOSURE_SHEETS) == {"D2-disc-listed", "D2-disc-soe"}
        with pytest.raises(HTTPException):
            resolve_variant("D2-2")


class TestRoundTrip:
    def test_listed_roundtrip_preserves_rows_and_notes(self):
        data = _sample_data()
        wb = _reload(build_disclosure_workbook("listed", SEGMENTS, data))
        storage, notes, warnings = parse_disclosure_workbook(wb, "listed", SEGMENTS)

        ind = storage["individual-rows"]
        assert len(ind) == 1
        assert ind[0]["name"] == "甲公司"
        assert ind[0]["endAmount"] == 100.0
        assert ind[0]["lossRate"] == 10.0
        assert ind[0]["basis"] == "已进入诉讼"
        assert "rowId" in ind[0]

        rev = storage["reversal-rows"][0]
        assert rev["cumulativeProvision"] == 20.0
        assert rev["recoveryMethod"] == "银行转账"

        wo = storage["writeoff-rows"][0]
        assert wo["amount"] == 15.0
        assert wo["relatedParty"] == "否"

        t5 = storage["top5-rows"][0]
        assert (t5["arAmount"], t5["contractAssetAmount"], t5["provision"]) == (1000.0, 200.0, 30.0)

        # 组合分表：父子结构还原，账龄 label → 段 key
        pf = storage["portfolios"]
        assert len(pf) == 1
        assert pf[0]["name"] == "应收中央企业客户"
        assert [r["key"] for r in pf[0]["rows"]] == ["within1", "y1to2"]
        assert pf[0]["rows"][0]["endAmount"] == 500.0

        assert notes["aging"] == "账龄结构说明"
        assert notes["top5"] == "前五名说明"
        # 未填写的子节不写入（不覆盖既有）
        assert "movement" not in notes
        assert warnings == []

        # 上市变体不解析国企专有表
        assert "other-portfolio-rows" not in storage
        assert "derecognized-rows" not in storage

    def test_soe_roundtrip_includes_soe_only_tables(self):
        wb = _reload(build_disclosure_workbook("soe", SEGMENTS, _sample_data()))
        storage, _notes, _warnings = parse_disclosure_workbook(wb, "soe", SEGMENTS)
        assert storage["other-portfolio-rows"][0]["name"] == "余额百分比组合"
        assert storage["derecognized-rows"][0]["gainLoss"] == -5.0


class TestNoWipeSemantics:
    def test_blank_template_import_writes_nothing(self):
        wb = _reload(build_disclosure_workbook("soe", SEGMENTS, None))
        storage, notes, warnings = parse_disclosure_workbook(wb, "soe", SEGMENTS)
        assert storage == {}
        assert notes == {}
        # 每张空表都提示"未改动既有数据"
        assert any("未改动既有数据" in w for w in warnings)

    def test_empty_note_does_not_overwrite(self):
        data = _sample_data()
        data["notes"] = {"aging": "   "}
        wb = _reload(build_disclosure_workbook("listed", SEGMENTS, data))
        _storage, notes, _warnings = parse_disclosure_workbook(wb, "listed", SEGMENTS)
        assert "aging" not in notes


class TestWarnings:
    def test_unmatched_aging_label_skipped_with_warning(self):
        data = _sample_data()
        data["portfolios"] = [
            {
                "groupId": "g1",
                "name": "组合A",
                "rows": [
                    {"key": "within1", "label": "1年以内", "endAmount": 10.0, "priorAmount": 5.0},
                    {"key": "legacy", "label": "9至10年", "endAmount": 99.0, "priorAmount": 0.0},
                ],
            }
        ]
        wb = _reload(build_disclosure_workbook("listed", SEGMENTS, data))
        storage, _notes, warnings = parse_disclosure_workbook(wb, "listed", SEGMENTS)
        assert [r["key"] for r in storage["portfolios"][0]["rows"]] == ["within1"]
        assert any("9至10年" in w for w in warnings)

    def test_aging_label_normalization_matches_variants(self):
        """「一年以内」「1～2年」等写法应归一匹配到当前段。"""
        wb = build_disclosure_workbook("listed", SEGMENTS, None)
        ws = wb[WS_PORTFOLIO]
        ws.append(["组合B", "一年以内", 1, 2])
        ws.append(["组合B", "1～2年", 3, 4])
        storage, _notes, warnings = parse_disclosure_workbook(_reload(wb), "listed", SEGMENTS)
        keys = [r["key"] for r in storage["portfolios"][0]["rows"]]
        assert keys == ["within1", "y1to2"]
        assert not any("未匹配" in w for w in warnings)

    def test_unknown_note_section_skipped_with_warning(self):
        wb = build_disclosure_workbook("listed", SEGMENTS, None)
        wb[WS_NOTES].append(["不存在的子节", "随便写"])
        _storage, notes, warnings = parse_disclosure_workbook(_reload(wb), "listed", SEGMENTS)
        assert notes == {}
        assert any("未识别的子节名" in w for w in warnings)

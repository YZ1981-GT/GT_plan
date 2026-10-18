"""回填只比对交付后改过的块（第 27 组，disclosure-note-doc-upload-and-batch-fill）。

现象（真实项目只读核对，2026-10-16）：交付中心附注默认走模板模式，块里除正文外还有模板的编号行、标题、指引段落
⇒ 块文字与库里正文本就不同。回填逐块拿块文字比库里正文：刚导出、未改一字的交付件，安徽 146 / 146 块、甲 90 / 90 块
全部判为已修改；库里正文自导出后没变 ⇒ 无冲突 ⇒ 全部写回（编号行、标题、指引并进正文，格式拍平）。
2026-08-04 回填验收探针在华氏写了 244 章次（之后已恢复）。

判据：当前块哈希 == 生成 / 刷新时记录的 Rendered_Block_Hash（与增量刷新的人工编辑检测同一判据）⇒ 没改过 ⇒
不比对、不写回；改过的块照旧走护栏、冲突检测与写回；记录为 NULL（存量交付件）的块照旧比对。

🔴 **附注回填已下线**（spec deliverable-writeback-and-online-edit-integrity，2026-10-17）：
生产 `_resolve_adapter` 对 disclosure_notes 返回 None、`writeback()` 短路返回全空。但「交付后
没改过的块不写回」是**与上游种类无关的通用机制**，仍是 audit_report 回填的活代码，须继续守护。
故本文件用例统一经 `_note_writeback_service` **强制注入附注适配器**（绕过 doc_type 门控），
验证块哈希判据本身；附注的「不可回填」由 test_deliverable_writeback_rowcount 的 decommission
用例守护。两者正交。
"""
from __future__ import annotations

from io import BytesIO

import sqlalchemy as sa
from docx import Document

from app.models.audit_platform_models import DeliverableSectionState
from app.models.report_models import DisclosureNote
from app.services.section_anchor_utils import (
    SectionBlock, block_text_hash, resolve_section_blocks, write_section_anchors,
)
from app.services.word_doc_utils import remove_section_markers, scan_section_blocks
from tests.test_deliverable_writeback_rowcount import (
    ACTOR_ID,
    PROJECT_ID,
    TASK_ID,
    YEAR,
    _note_writeback_service,
    _run,
    _setup,
)

#: 模板模式交付件的块形态：编号行 + 标题 + 正文 + 模板指引（正文之外的都不在库里）
BLOCKS = {
    "八、1": ["1", "货币资金", "货币资金正文", "（注：按币种披露。）"],
    "八、2": ["2", "应收票据", "应收票据正文"],
}
NOTES = {"八、1": "货币资金正文", "八、2": "应收票据正文"}


def _delivered(blocks: dict[str, list[str]]) -> bytes:
    doc = Document()
    for code, paragraphs in blocks.items():
        doc.add_paragraph(f"##SECTION:{code}##")
        for text in paragraphs:
            doc.add_paragraph(text)
        doc.add_paragraph(f"##/SECTION:{code}##")
    scanned = {b.section_code: b for b in scan_section_blocks(doc)}
    write_section_anchors(doc, [SectionBlock(section_code=c, open_el=scanned[c].open_el, close_el=scanned[c].close_el)
                                for c in blocks])
    remove_section_markers(doc)
    buf = BytesIO()
    doc.save(buf)
    return buf.getvalue()


def _edited(data: bytes, old: str, new: str) -> bytes:
    doc = Document(BytesIO(data))
    [para] = [p for p in doc.paragraphs if p.text == old]
    para.text = new
    buf = BytesIO()
    doc.save(buf)
    return buf.getvalue()


async def _record_rendered_hashes(session, data: bytes) -> None:
    mode, blocks = resolve_section_blocks(Document(BytesIO(data)))
    assert mode == "anchor"
    for block in blocks:
        await session.execute(sa.update(DeliverableSectionState)
                              .where(DeliverableSectionState.section_code == block.section_code)
                              .values(rendered_block_hash=block_text_hash(list(block.elements))))
    await session.flush()


async def _texts(session) -> dict[str, str]:
    return dict((await session.execute(sa.select(DisclosureNote.note_section, DisclosureNote.text_content))).all())


def test_untouched_template_blocks_are_not_written_back():
    async def scenario(session):
        await _setup(session, notes=NOTES, states=list(NOTES))
        delivered = _delivered(BLOCKS)
        await _record_rendered_hashes(session, delivered)
        result = await _note_writeback_service(session).writeback(
            TASK_ID, PROJECT_ID, YEAR, ACTOR_ID, docx_bytes=delivered)
        return result, await _texts(session)

    result, texts = _run(scenario)
    assert (result["written"], result["conflicts"], result["failed"]) == ([], [], [])
    assert texts == NOTES, "没改过的块把编号 / 标题 / 指引写进了库里正文"


def test_edited_block_is_still_compared_and_written():
    async def scenario(session):
        await _setup(session, notes=NOTES, states=list(NOTES))
        delivered = _delivered(BLOCKS)
        await _record_rendered_hashes(session, delivered)
        result = await _note_writeback_service(session).writeback(
            TASK_ID, PROJECT_ID, YEAR, ACTOR_ID, docx_bytes=_edited(delivered, "货币资金正文", "货币资金正文（已改）"))
        return result, await _texts(session)

    result, texts = _run(scenario)
    assert result["written"] == ["八、1"]
    assert "货币资金正文（已改）" in texts["八、1"]
    assert texts["八、2"] == NOTES["八、2"], "没改过的兄弟块被一并写回"


def test_blocks_without_recorded_hash_keep_legacy_comparison():
    """对照：存量交付件（Rendered_Block_Hash 为 NULL）照旧逐块比对 —— 证明上面的不写回来自哈希判据。"""
    async def scenario(session):
        await _setup(session, notes=NOTES, states=list(NOTES))
        return await _note_writeback_service(session).writeback(
            TASK_ID, PROJECT_ID, YEAR, ACTOR_ID, docx_bytes=_delivered(BLOCKS))

    result = _run(scenario)
    assert sorted(result["written"]) == ["八、1", "八、2"]

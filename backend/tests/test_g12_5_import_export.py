"""G12-5 五区段导入导出单元测试."""

from __future__ import annotations

from io import BytesIO

from app.routers.wp_render_strategies._g12_net_hedge_gains_import_export import (
    _G12_5_SHEET_META,
    _G12_5_SHEET_TEST,
    _build_g12_5_workbook,
    _parse_g12_5_import,
)


def test_g12_5_workbook_template_has_meta_sheet():
    wb = _build_g12_5_workbook(None, template_only=True)
    assert _G12_5_SHEET_TEST in wb.sheetnames
    assert _G12_5_SHEET_META in wb.sheetnames
    assert "编制说明" in wb.sheetnames


def test_g12_5_export_import_roundtrip():
    payload = {
        "rows": [{
            "seq": 1,
            "hedgeRelationId": "HR-1",
            "item": "外汇净头寸",
            "currency": "USD",
            "position1Desc": "销售",
            "position1Amount": "1000万美元",
            "position2Desc": "采购",
            "position2Amount": "1200万美元",
            "netPosition": "支付200万美元",
            "evidenceType": "sales_budget",
            "supportingEvidence": "预算",
            "hedgingInstrument": "远期",
            "indexRef": "G12-2/1",
        }],
        "testObjective": "检查净敞口",
        "sampleCriteria": "全部运用套期的净敞口",
        "auditNote": "已获取支持性证据",
        "conclusion": "未发现异常",
    }
    wb = _build_g12_5_workbook(payload, template_only=False)
    buf = BytesIO()
    wb.save(buf)
    parsed, errors, count = _parse_g12_5_import(buf.getvalue())
    assert count == 1
    assert not errors
    assert parsed["testObjective"] == "检查净敞口"
    assert parsed["sampleCriteria"] == "全部运用套期的净敞口"
    assert parsed["auditNote"] == "已获取支持性证据"
    assert parsed["conclusion"] == "未发现异常"
    assert parsed["rows"][0]["item"] == "外汇净头寸"
    assert parsed["rows"][0]["currency"] == "USD"
    assert parsed["rows"][0]["evidenceType"] == "sales_budget"


def test_g12_5_import_legacy_single_sheet_only():
    """兼容仅含 G12-5 测试行的旧版导出."""
    wb = _build_g12_5_workbook([{
        "seq": 1,
        "item": "仅测试区",
        "currency": "CNY",
        "position1Desc": "A",
        "position2Desc": "B",
        "netPosition": "零",
    }], template_only=False)
    del wb[_G12_5_SHEET_META]
    buf = BytesIO()
    wb.save(buf)
    parsed, errors, count = _parse_g12_5_import(buf.getvalue())
    assert count == 1
    assert parsed["rows"][0]["item"] == "仅测试区"


def test_g12_5_import_legacy_without_evidence_type():
    """旧版无「证据类型」列仍可导入."""
    from openpyxl import Workbook

    wb = Workbook()
    ws = wb.active
    ws.title = "G12-5"
    ws.append(["G12-5 风险净敞口检查表"])
    ws.append([
        "序号", "套期关系编号", "项目", "币种", "头寸1描述", "头寸1金额", "头寸2描述", "头寸2金额",
        "净头寸", "支持性证据", "套期工具", "索引号",
    ])
    ws.append([1, "HR-1", "旧行", "USD", "A", "100", "B", "80", "收20", "销售预算", "远期", "G12-2"])
    buf = BytesIO()
    wb.save(buf)
    parsed, errors, count = _parse_g12_5_import(buf.getvalue())
    assert count == 1
    assert parsed["rows"][0]["item"] == "旧行"
    assert parsed["rows"][0]["supportingEvidence"] == "销售预算"
    assert parsed["rows"][0].get("evidenceType", "") == ""


def test_g12_5_import_empty_raises_error():
    wb = _build_g12_5_workbook(None, template_only=True)
    buf = BytesIO()
    wb.save(buf)
    _, errors, count = _parse_g12_5_import(buf.getvalue())
    assert count == 0
    assert any("未解析" in e for e in errors)

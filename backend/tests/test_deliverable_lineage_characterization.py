"""交付 docx 灰度关闭 characterization 基线 — deliverable-lineage-content-control Task 1（子任务 4）.

目的
----
为「灰度开关 ``DELIVERABLE_LINEAGE_CONTENT_CONTROL_ENABLED`` 关闭时生成的 docx」冻结一份
**确定性结构快照**，供 Task 3.1 的 Property 3（逐字节/结构等价对比）作为参照基线，
防止后续改动在开关关闭时意外泄漏内容控件或改变章节可见内容/顺序（零回归红线）。

确定性保证（避免 flaky）
------------------------
- 快照取 **body 级结构指纹**（元素类型 + 规范化文本 + 有序 bookmark 名 + 有序 sdt Tag），
  不比较 zip 原始字节，因此不受 docx 打包时间戳/随机字节影响。
- 合成 docx 由固定内容构造（无时间、无随机、无 DB），bookmark id 显式指定。
- 模拟 note 导出器 step6→step7 之间的收尾序列（``if flag: 注入`` → ``remove_section_markers``），
  与 ``note_word_exporter.export`` 的灰度守卫逐字对应。

不依赖 DB / OnlyOffice。
"""

from __future__ import annotations

from docx import Document
from docx.oxml import OxmlElement
from docx.oxml.ns import qn

from app.core.config import settings
from app.services.content_control_injector import inject_content_controls_for_blocks
from app.services.section_anchor_utils import anchor_name
from app.services.word_doc_utils import remove_section_markers, scan_section_blocks


# ---------------------------------------------------------------------------
# 合成 note 导出中间态 docx（step6 之后、step7 之前的状态）
# ---------------------------------------------------------------------------


def _add_bookmark(paragraph, name: str, bm_id: str) -> None:
    """在段落内加一个 note_section bookmark（模拟 note 导出书签，验证并存）。"""
    start = OxmlElement("w:bookmarkStart")
    start.set(qn("w:id"), bm_id)
    start.set(qn("w:name"), name)
    paragraph._p.insert(0, start)
    end = OxmlElement("w:bookmarkEnd")
    end.set(qn("w:id"), bm_id)
    paragraph._p.append(end)


def _build_export_intermediate_doc():
    """构造与 note 导出 step6 后一致的 SECTION 块中间态（含 note_section bookmark）。"""
    doc = Document()
    doc.add_paragraph("##SECTION:八、1##")
    p_a = doc.add_paragraph("货币资金内容 A")
    doc.add_paragraph("货币资金内容 B")
    doc.add_paragraph("##/SECTION:八、1##")
    doc.add_paragraph("##SECTION:五、1##")
    doc.add_paragraph("应收账款内容")
    doc.add_paragraph("##/SECTION:五、1##")
    _add_bookmark(p_a, "note_section_八、1", "5001")
    return doc


# ---------------------------------------------------------------------------
# 收尾序列（严格镜像 note_word_exporter.export 的 step6.5 灰度守卫 + step7）
# ---------------------------------------------------------------------------


def _finalize(doc, *, flag: bool) -> None:
    """镜像 note 导出收尾：``if flag: 注入内容控件`` → 清理 SECTION 标记。"""
    if flag:
        inject_content_controls_for_blocks(doc, scan_section_blocks(doc))
    remove_section_markers(doc)


# ---------------------------------------------------------------------------
# 确定性结构指纹
# ---------------------------------------------------------------------------


def _para_text(el) -> str:
    return "".join(t.text or "" for t in el.iter(qn("w:t"))).strip()


def structural_fingerprint(doc) -> dict:
    """body 级确定性结构指纹（无时间戳/随机字节）。

    - ``blocks``：body 直接子元素按序 [(kind, text)]（p 带规范化文本 / tbl 空文本）
    - ``bookmarks``：有序 bookmark 名列表
    - ``sdt_tags``：有序内容控件 Tag 列表
    """
    body = doc.element.body
    blocks: list[tuple[str, str]] = []
    for child in body.iterchildren():
        if child.tag == qn("w:p"):
            blocks.append(("p", _para_text(child)))
        elif child.tag == qn("w:tbl"):
            blocks.append(("tbl", ""))
        elif child.tag == qn("w:sdt"):
            blocks.append(("sdt", ""))
    bookmarks = [
        bm.get(qn("w:name")) for bm in body.iter(qn("w:bookmarkStart"))
    ]
    sdt_tags: list[str] = []
    for sdt in body.iter(qn("w:sdt")):
        pr = sdt.find(qn("w:sdtPr"))
        tag_el = pr.find(qn("w:tag")) if pr is not None else None
        if tag_el is not None:
            sdt_tags.append(tag_el.get(qn("w:val")))
    return {"blocks": blocks, "bookmarks": bookmarks, "sdt_tags": sdt_tags}


def visible_texts(doc) -> list[str]:
    """全文可见段落文本（含 sdt 内下钻）——顺序敏感。"""
    return [_para_text(p) for p in doc.element.body.iter(qn("w:p"))]


# ---------------------------------------------------------------------------
# 冻结基线（灰度关闭时的确定性结构快照）
# ---------------------------------------------------------------------------

# note 导出收尾（flag=False）后：SECTION 开闭标记段落被 remove_section_markers 删除，
# 内容段落原样保留、无任何 w:sdt、note_section bookmark 保留。
BASELINE_FLAG_OFF = {
    "blocks": [
        ("p", "货币资金内容 A"),
        ("p", "货币资金内容 B"),
        ("p", "应收账款内容"),
    ],
    "bookmarks": ["note_section_八、1"],
    "sdt_tags": [],
}


# ---------------------------------------------------------------------------
# 测试
# ---------------------------------------------------------------------------


def test_flag_default_is_false():
    """灰度开关缺省 = False（Requirement 6.1 / Property 11）。"""
    assert settings.DELIVERABLE_LINEAGE_CONTENT_CONTROL_ENABLED is False


def test_flag_off_matches_frozen_baseline():
    """灰度关闭：收尾后结构指纹 == 冻结基线，且零 w:sdt（characterization 快照）。"""
    doc = _build_export_intermediate_doc()
    _finalize(doc, flag=False)
    fp = structural_fingerprint(doc)
    assert fp == BASELINE_FLAG_OFF
    assert fp["sdt_tags"] == []  # 关闭时绝不注入内容控件


def test_flag_off_snapshot_is_deterministic():
    """同输入两次收尾（flag=False）结构指纹完全一致（无时间戳/随机字节 → 不 flaky）。"""
    fp1 = structural_fingerprint(_finalized(flag=False))
    fp2 = structural_fingerprint(_finalized(flag=False))
    assert fp1 == fp2 == BASELINE_FLAG_OFF


def test_flag_off_equals_pre_feature_path():
    """Property 3：灰度关闭路径与「引入本能力前」（纯 remove_section_markers，无任何注入）结构等价。"""
    # 引入前：不存在内容控件注入，仅清理标记
    pre = _build_export_intermediate_doc()
    remove_section_markers(pre)
    # 灰度关闭：if flag(False) 跳过注入 → 清理标记
    off = _build_export_intermediate_doc()
    _finalize(off, flag=False)
    assert structural_fingerprint(pre) == structural_fingerprint(off)


def test_flag_on_only_adds_sdt_wrapper():
    """灰度开启：唯一差异是新增 sec_ 内容控件包裹；可见文本序列与 bookmark 与关闭态一致。"""
    off = _finalized(flag=False)
    on = _finalized(flag=True)

    fp_off = structural_fingerprint(off)
    fp_on = structural_fingerprint(on)

    # 开启态注入了 sec_ Tag 内容控件
    assert anchor_name("八、1") in fp_on["sdt_tags"]
    assert anchor_name("五、1") in fp_on["sdt_tags"]
    assert fp_off["sdt_tags"] == []

    # 可见文本（下钻 sdt 内）与 bookmark 集合不变——仅结构外层多了 w:sdt 包裹
    assert visible_texts(on) == visible_texts(off)
    assert set(fp_on["bookmarks"]) == set(fp_off["bookmarks"])


def _finalized(*, flag: bool):
    doc = _build_export_intermediate_doc()
    _finalize(doc, flag=flag)
    return doc

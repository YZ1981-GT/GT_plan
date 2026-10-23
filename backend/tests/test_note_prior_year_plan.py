"""上年附注 Word → 本年附注表格期初列：解析 / 逐表规划 / 写入（spec consol-note-header-library-and-node-reaggregate
R6.6，design §15.6）。

原实现只读 Word 表第一行当表头、按原始单元格位置取值、按表序 zip 配对、跳过内容控件里的内容 ⇒ 分组表找不到列或填错列、
平台导出的交付件整份解析为空。这里用真 docx（python-docx 生成，含合并单元格 / 内容控件）逐条钉住。
"""
from __future__ import annotations

import io
from typing import Any

import docx
import pytest
from docx.oxml import OxmlElement

from app.services.note_prior_year_plan import (
    STATUS_AUTO,
    STATUS_CONFIRM,
    STATUS_NONE,
    apply_col_map,
    match_sections,
    parse_amount,
    parse_prior_docx,
    plan_note_tables,
)


def _add_table(doc: Any, rows: list[list[str | None]], merges: tuple = ()) -> Any:
    t = doc.add_table(rows=len(rows), cols=len(rows[0]))
    for r0, c0, r1, c1 in merges:
        t.cell(r0, c0).merge(t.cell(r1, c1))
    for r, row in enumerate(rows):
        for c, value in enumerate(row):
            if value is not None:
                t.cell(r, c).text = value
    return t


def _save(doc: Any) -> bytes:
    buf = io.BytesIO()
    doc.save(buf)
    return buf.getvalue()


def _note(title: str, tables: list[dict], nid: str = "n1") -> dict:
    return {"id": nid, "note_section": "八、1", "section_title": title, "table_data": {"_tables": tables}}


def _rows(*labels: str, width: int) -> list[dict]:
    return [{"label": label, "values": [None] * width} for label in labels]


def _plan(data: bytes, note: dict):
    doc = parse_prior_docx(data)
    [h] = match_sections([note["section_title"]], doc)
    return doc, plan_note_tables(note, doc.headings[h] if h is not None else None, doc)


#: 两级表头：期末余额 / 期初余额 各含 账面余额 | 坏账准备（「项目」纵向合并）
_GROUPED = [
    ["项  目", "期末余额", None, "期初余额", None],
    [None, "账面余额", "坏账准备", "账面余额", "坏账准备"],
    ["1年以内", "1,000.00", "50", "800", "40"],
    ["合  计", "1,000.00", "50", "800", "40"],
]
_GROUPED_MERGES = ((0, 0, 1, 0), (0, 1, 0, 2), (0, 3, 0, 4))


def _grouped_table(name: str = "按账龄披露应收账款") -> dict:
    return {"name": name, "headers": ["账龄", "账面余额", "坏账准备", "账面余额", "坏账准备"],
            "_column_groups": [{"group": "期末余额", "start": 1, "span": 2}, {"group": "期初余额", "start": 3, "span": 2}],
            "rows": _rows("1年以内", "合计", width=4)}


class TestParse:
    def test_grouped_header_columns_and_values_follow_the_grid(self):
        """分组表：表头按网格取两级路径，数据行按逻辑列取值（原实现只读第一行当表头、按原始单元格位置取值）。"""
        doc = docx.Document()
        doc.add_heading("应收账款", level=2)
        _add_table(doc, _GROUPED, _GROUPED_MERGES)
        parsed = parse_prior_docx(_save(doc))
        [table] = parsed.tables.values()
        assert [k.label for k in table.keys[1:]] == ["期末余额/账面余额", "期末余额/坏账准备", "期初余额/账面余额", "期初余额/坏账准备"]
        assert [k.period for k in table.keys[1:]] == ["end", "end", "prior", "prior"]
        assert table.word.rows == [["1年以内", "1,000.00", "50", "800", "40"], ["合  计", "1,000.00", "50", "800", "40"]]
        assert parsed.headings[0].tables == [table.id]

    def test_content_controls_are_read(self):
        """平台导出的交付件每节都包在内容控件 w:sdt 里：原实现只看 body 直接子元素 ⇒ 整份解析为空。"""
        doc = docx.Document()
        head = doc.add_heading("货币资金", level=2)._p
        tbl = _add_table(doc, [["项目", "期末余额", "期初余额"], ["库存现金", "12", "3"]])._tbl
        sdt, content = OxmlElement("w:sdt"), OxmlElement("w:sdtContent")
        sdt.append(content)
        content.append(head)
        content.append(tbl)
        doc.element.body.insert(0, sdt)
        parsed = parse_prior_docx(_save(doc))
        assert [h.title for h in parsed.headings] == ["货币资金"]
        assert [t.word.rows for t in parsed.tables.values()] == [[["库存现金", "12", "3"]]]

    def test_table_names_before_explanatory_text_are_captions(self):
        """表名与表之间隔着说明文字：表前的全部段落都是表名候选；小标题下的表属于上一级节。"""
        doc = docx.Document()
        doc.add_heading("应收账款", level=2)
        doc.add_heading("（1）按账龄披露应收账款", level=3)
        doc.add_paragraph("账龄自应收账款确认日起算。")
        _add_table(doc, _GROUPED, _GROUPED_MERGES)
        parsed = parse_prior_docx(_save(doc))
        [table] = parsed.tables.values()
        assert table.captions == ["（1）按账龄披露应收账款", "账龄自应收账款确认日起算。"]
        assert [h.tables for h in parsed.headings] == [[table.id], [table.id]]

    def test_split_cells_in_a_data_row(self):
        """数据行把一列拆成两格（表头只有一格）：取该逻辑列里从本行起、有字的第一格。"""
        doc = docx.Document()
        doc.add_heading("货币资金", level=2)
        t = _add_table(doc, [["项目", "期末余额", "期末余额", "期初余额"], ["库存现金", "", "12", "3"]],
                       merges=((0, 1, 0, 2),))
        assert t is not None
        parsed = parse_prior_docx(_save(doc))
        [table] = parsed.tables.values()
        assert table.word.rows == [["库存现金", "12", "3"]]


class TestPlan:
    def test_grouped_opening_columns_take_prior_closing_columns(self):
        doc = docx.Document()
        doc.add_heading("应收账款", level=2)
        doc.add_paragraph("（1）按账龄披露应收账款")
        _add_table(doc, _GROUPED, _GROUPED_MERGES)
        table = _grouped_table()
        _doc, [plan] = _plan(_save(doc), _note("应收账款", [table]))
        assert (plan.status, plan.method, plan.col_map) == (STATUS_AUTO, "name", {3: 1, 4: 2})
        assert apply_col_map(table, _doc.tables[plan.prior_table_id].word, plan.col_map) == 4
        assert [r["values"] for r in table["rows"]] == [[None, None, 1000.0, 50.0], [None, None, 1000.0, 50.0]]

    def test_continuation_takes_prior_main_table(self):
        """本年「（续：期初数）」← 上年「（期末数）」主表的期末列；上年续表是前年的数，不取。"""
        doc = docx.Document()
        doc.add_heading("应收票据", level=2)
        doc.add_paragraph("（1）按坏账计提方法分类（期末数）")
        _add_table(doc, [["类别", "账面余额", "坏账准备"], ["按单项计提", "10", "1"]])
        doc.add_paragraph("按坏账计提方法分类（续：期初数）")
        _add_table(doc, [["类别", "账面余额", "坏账准备"], ["按单项计提", "7", "0.5"]])
        main = {"name": "按坏账计提方法分类（期末数）", "headers": ["类别", "账面余额", "坏账准备"], "rows": _rows("按单项计提", width=2)}
        cont = {"name": "按坏账计提方法分类（续：期初数）", "headers": ["类别", "账面余额", "坏账准备"],
                "rows": _rows("按单项计提", width=2)}
        parsed, plans = _plan(_save(doc), _note("应收票据", [main, cont]))
        assert [(p.table_index, p.status, p.prior_table_id, p.col_map) for p in plans] == [(1, STATUS_AUTO, "t1", {1: 1, 2: 2})]
        apply_col_map(cont, parsed.tables["t1"].word, plans[0].col_map)
        assert cont["rows"][0]["values"] == [10.0, 1.0]

    def test_main_table_without_period_takes_opposite_of_its_continuation(self):
        """上年主表表头没写期间、后面隔着别的表跟「（续：期初数）」列续表（交付件写法）：主表按续表取对侧 = 期末，
        本年续表（期初数）← 上年主表。原来两张同基名的表都对不上列 ⇒ 待确认且列对应为空（真库：和平药房 ← 重药安徽交付件）。"""
        doc = docx.Document()
        doc.add_heading("应收账款", level=2)
        doc.add_paragraph("（1）按坏账准备计提方法分类披露应收账款")
        _add_table(doc, [["类别", "账面余额", "坏账准备", "账面价值"], ["按组合计提", "100", "5", "95"]])
        doc.add_paragraph("（2）按欠款方归集的期末余额前五名")
        _add_table(doc, [["债务人名称", "账面余额"], ["甲医院", "30"]])
        doc.add_paragraph("（1）按坏账准备计提方法分类披露应收账款（续：期初数）")
        _add_table(doc, [["类别", "账面余额", "坏账准备", "账面价值"], ["按组合计提", "80", "4", "76"]])
        headers = ["类别", "账面余额", "坏账准备", "账面价值"]
        main = {"name": "按坏账准备计提方法分类披露应收账款", "headers": headers, "rows": _rows("按组合计提", width=3)}
        cont = {"name": "按坏账准备计提方法分类披露应收账款（续：期初数）", "headers": headers,
                "rows": _rows("按组合计提", width=3)}
        parsed, plans = _plan(_save(doc), _note("应收账款", [main, cont]))
        assert [t.continuations for t in parsed.tables.values()] == [["t3"], [], []]
        assert [(p.table_index, p.status, p.method, p.prior_table_id, p.col_map) for p in plans] == [
            (1, STATUS_AUTO, "name", "t1", {1: 1, 2: 2, 3: 3})]
        apply_col_map(cont, parsed.tables["t1"].word, plans[0].col_map)
        assert cont["rows"][0]["values"] == [100.0, 5.0, 95.0]

    def test_continuation_is_not_linked_to_another_section_by_adjacency_alone(self):
        """续表只凭「紧挨着上一张」认主表时须在同一标题下：交付件删掉了空主表，续表不认到上一节的表上去补期间。"""
        doc = docx.Document()
        doc.add_heading("货币资金", level=2)
        _add_table(doc, [["项目", "账面余额"], ["甲", "5"]])
        doc.add_heading("应收票据", level=2)
        doc.add_paragraph("应收票据分类（续：期初数）")
        _add_table(doc, [["项目", "账面余额"], ["甲", "4"]])
        assert [t.continuations for t in parse_prior_docx(_save(doc)).tables.values()] == [[], []]

    def test_long_explanatory_paragraph_is_not_shown_as_the_table_name(self):
        """平台导出的表前常是一段正文：展示用表名取离表最近的短段落；只有长段说明时取首行截断（原样塞进原因里一大段）。"""
        doc = docx.Document()
        doc.add_heading("货币资金", level=2)
        doc.add_paragraph("其他货币资金明细")
        doc.add_paragraph("其他货币资金主要包含银行承兑汇票保证金、信用证保证金、外埠存款等，期末余额为若干元，具体见下表。")
        _add_table(doc, [["项目", "期末余额"], ["甲", "5"]])
        doc.add_paragraph("说明：" + "很长的说明文字" * 10)
        _add_table(doc, [["项目", "期末余额"], ["乙", "6"]])
        first, second = parse_prior_docx(_save(doc)).tables.values()
        assert first.caption == "其他货币资金明细"
        assert second.caption.endswith("…") and len(second.caption) == 41

    def test_listed_prior_year_end_column(self):
        """上市「上年年末余额」是期初一侧：取上年表的期末余额。"""
        doc = docx.Document()
        doc.add_heading("货币资金", level=2)
        _add_table(doc, [["项目", "期末余额", "上年年末余额"], ["库存现金", "12", "3"]])
        table = {"name": "货币资金", "headers": ["项目", "期末余额", "上年年末余额"], "rows": _rows("库存现金", width=2)}
        _doc, [plan] = _plan(_save(doc), _note("货币资金", [table]))
        assert (plan.status, plan.col_map) == (STATUS_AUTO, {2: 1})

    def test_no_name_evidence(self):
        """没有表名证据：只有一张表 ⇒ 用它；两年表数相同 ⇒ 同位置的表；否则按结构择优要确认 / 对不上就不选。"""
        def build(n: int) -> bytes:
            doc = docx.Document()
            doc.add_heading("应收账款", level=2)
            for _ in range(n):
                _add_table(doc, [["项目", "期末余额", "期初余额"], ["甲", "5", "4"]])
            return _save(doc)

        def table(name: str) -> dict:
            return {"name": name, "headers": ["项目", "期末余额", "期初余额"], "rows": _rows("甲", width=2)}

        _d, [single] = _plan(build(1), _note("应收账款", [table("无关表名")]))
        assert (single.status, single.method) == (STATUS_AUTO, "single_table")
        _d, plans = _plan(build(2), _note("应收账款", [table("a"), table("b")]))
        assert [(p.status, p.method, p.prior_table_id) for p in plans] == [
            (STATUS_AUTO, "position", "t1"), (STATUS_AUTO, "position", "t2")]
        _d, plans = _plan(build(3), _note("应收账款", [table("a"), table("b")]))
        assert {p.status for p in plans} == {STATUS_CONFIRM} and all(p.prior_table_id is None for p in plans)
        assert "上年该节有 3 张表" in plans[0].reason and plans[0].candidates == ["t1", "t2", "t3"]

    def test_structure_choice_needs_confirmation(self):
        doc = docx.Document()
        doc.add_heading("应收账款", level=2)
        _add_table(doc, [["项目", "期末余额", "期初余额"], ["甲", "5", "4"]])
        _add_table(doc, [["项目", "本期增加", "本期减少"], ["甲", "1", "2"]])
        table = {"name": "x", "headers": ["项目", "期末余额", "期初余额"], "rows": _rows("甲", width=2)}
        other = {"name": "y", "headers": ["项目", "期末余额", "期初余额", "期初数量"], "rows": _rows("甲", width=3)}
        _d, plans = _plan(_save(doc), _note("应收账款", [table, other, {"name": "z", "headers": ["项目", "期末余额"], "rows": []}]))
        assert [(p.status, p.method, p.prior_table_id) for p in plans] == [
            (STATUS_CONFIRM, "structure", "t1"), (STATUS_CONFIRM, "structure", "t1")]
        assert "按表头选了" in plans[0].reason
        assert "期初列「期初数量」在上年表" in plans[1].reason and plans[1].col_map == {2: 1}

    def test_no_source(self):
        doc = docx.Document()
        doc.add_heading("货币资金", level=2)
        doc.add_paragraph("没有表格")
        table = {"name": "存货", "headers": ["项目", "期末余额", "期初余额"], "rows": _rows("原材料", width=2)}
        _d, [missing] = _plan(_save(doc), _note("存货", [table]))
        assert (missing.status, missing.reason) == (STATUS_NONE, "上年附注里没有找到「存货」这一节")
        _d, [empty] = _plan(_save(doc), _note("货币资金", [table]))
        assert (empty.status, empty.reason) == (STATUS_NONE, "上年附注「货币资金」这一节没有表格")

    def test_same_titles_pair_in_order(self):
        """同名章节（合并 / 母公司「应收账款」）按出现序号对应。"""
        doc = docx.Document()
        for title in ("应收账款", "其他", "应收账款"):
            doc.add_heading(title, level=2)
        parsed = parse_prior_docx(_save(doc))
        assert match_sections(["应收账款", "应收账款", "存货"], parsed) == [0, 2, None]


class TestPlatformExportRoundTrip:
    @pytest.mark.asyncio
    async def test_editor_word_export_can_be_imported_next_year(self):
        """编辑器「导出 Word」原来标题全是正文样式（只加粗 / 改字号），目录域与导航窗格是空的，上年导入 / 上传附注文档
        一个标题都认不出（真库：重药安徽导出 → 和平药房预览 153 张表全部「上年没有」）。标题设大纲级别后往返可认、可填。"""
        from types import SimpleNamespace
        from unittest.mock import AsyncMock
        from uuid import uuid4

        from docx.oxml.ns import qn

        from app.services.note_word_exporter import NoteWordExporter

        def note(section: str, title: str, text: str, table: dict | None) -> SimpleNamespace:
            return SimpleNamespace(note_section=section, section_title=title, text_content=text, table_data=table,
                                   guidance_text=None, is_deleted=False, status="draft", is_empty=False,
                                   source_template=None)

        prior = {"name": "货币资金", "headers": ["项目", "期末余额", "期初余额"],
                 "rows": [{"label": "库存现金", "values": [12, 3]}]}
        exporter = NoteWordExporter(SimpleNamespace())  # type: ignore[arg-type]
        exporter._load_notes = AsyncMock(return_value=[  # type: ignore[method-assign]
            note("八", "项目注释", "以下注释项目除特别注明外，金额单位为人民币元。", None),
            note("八、1", "货币资金", "货币资金说明。", prior)])
        data = (await exporter.export(uuid4(), 2025, template_type="soe", report_scope="standalone")).getvalue()
        document = docx.Document(io.BytesIO(data))
        outline = {p.text: p._p.pPr.find(qn("w:outlineLvl")).get(qn("w:val"))
                   for p in document.paragraphs if p._p.pPr is not None and p._p.pPr.find(qn("w:outlineLvl")) is not None}
        assert outline == {"一、项目注释": "0", "（一）货币资金": "1"}
        current = {"name": "货币资金", "headers": ["项目", "期末余额", "期初余额"], "rows": _rows("库存现金", width=2)}
        parsed, [plan] = _plan(data, _note("货币资金", [current]))
        assert [(h.level, h.title) for h in parsed.headings] == [(1, "一、项目注释"), (2, "（一）货币资金")]
        assert (plan.status, plan.col_map) == (STATUS_AUTO, {2: 1})
        assert apply_col_map(current, parsed.tables[plan.prior_table_id].word, plan.col_map) == 1
        assert current["rows"][0]["values"] == [None, 12.0]


class TestApply:
    def test_rows_by_label_and_ordinal_and_only_empty_cells(self):
        doc = docx.Document()
        doc.add_heading("其他应付款", level=2)
        _add_table(doc, [["项目", "期末余额", "期初余额"], ["其他", "1", "9"], ["其他", "(2)", "9"], ["合  计：", "-", "9"]])
        table = {"name": "其他应付款", "headers": ["项目", "期末余额", "期初余额"],
                 "rows": [{"label": "其他", "values": [None, None]}, {"label": "其他", "values": [None, 8]},
                          {"label": "……", "values": [None, None]}, {"label": "合计", "values": [None, None]}]}
        parsed, [plan] = _plan(_save(doc), _note("其他应付款", [table]))
        assert apply_col_map(table, parsed.tables[plan.prior_table_id].word, plan.col_map) == 1
        assert [r["values"] for r in table["rows"]] == [[None, 1.0], [None, 8], [None, None], [None, None]]

    @pytest.mark.parametrize(("text", "value"), [
        ("1,234.50", 1234.5), ("(66)", -66.0), ("（12.5）", -12.5), ("-", None), ("—", None), ("", None), ("X.XX", None),
    ])
    def test_parse_amount(self, text, value):
        assert parse_amount(text) == value

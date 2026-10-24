"""三参 PREV 预填逐格定位表（spec formula-fetch-followup-defect-closure R2）。

修复前：初次预填只读每个底稿的第一块映射，三参 PREV 没有取值分支；位置是语义名，按列头关键词猜格
（「上年审定数」= 列头含「上年」或「审定数」+ 第 7 行起第一个数据行）⇒ 落到「期末与期初审定数比较」「本期审定数」
列或分节标题行；重新预填叠加有效公式时同一套猜法。

守卫（真模板）：映射全部三参格 ↔ 定位表一一对应；目标格签名（列头 / 行标签 / 公式标记）与系统模板一致、动作与公式
标记一致、目标格属上年 / 期初语义；同一目标格不许两套取数参数；生成器输出与入库文件逐字一致（``--check``），模板
列头一改就红。运行时（模板副本 + SQLite 真 ORM）：``write`` 写值 / ``verify`` 只批注 / 取不到批注原因不写 0 /
模板与定位表不符不动 / 用户公式格不动；初次预填不再按关键词猜三参格；重新预填 ``locate_cell`` 只认定位表；
路由端到端（上年项目 → 批注上年值）；试算表取数按审计年度且排除已删除行。
"""
from __future__ import annotations

import hashlib
import importlib.util
import shutil
import uuid
from decimal import Decimal
from pathlib import Path

import pytest
from openpyxl import load_workbook

from app.services import prev3_prefill_locator as prev3

BACKEND = Path(__file__).resolve().parents[1]
_GEN = BACKEND / "scripts" / "gen" / "gen_prev3_prefill_locator.py"


def _load_gen():
    spec = importlib.util.spec_from_file_location("gen_prev3_prefill_locator_t", _GEN)
    module = importlib.util.module_from_spec(spec)
    spec.loader.exec_module(module)  # type: ignore[union-attr]
    return module


@pytest.fixture(scope="module")
def gen():
    return _load_gen()


@pytest.fixture(scope="module")
def built(gen):
    """生成器现扫一次（全部系统模板，约 20 秒）。"""
    return gen.build()


@pytest.fixture(scope="module")
def mapping_cells(gen):
    import json

    return gen.prev3_cells(json.loads(gen.MAPPING.read_text(encoding="utf-8")))


def _entries() -> dict[tuple[str, str, str], prev3.LocatorEntry]:
    entries = prev3.load_locator()
    assert entries, "定位表为空或读不出来：后面的逐条断言会空转"
    return entries


# ── 定位表 ↔ 映射表 ↔ 模板 ─────────────────────────────────────────────────────────────


def test_locator_file_is_generator_output(gen, built):
    """入库的定位表 = 生成器现扫结果（等价 ``--check``）；且没有待人工裁定的格。"""
    doc, unresolved = built
    assert unresolved == [], f"自动规则定不了、又没有人工裁定的格：{unresolved}"
    assert gen.render(doc) == prev3.LOCATOR_FILE.read_text(encoding="utf-8"), "定位表与模板 / 映射表不同步，请重跑生成器"


def test_every_prev3_cell_has_exactly_one_entry(mapping_cells):
    entries = _entries()
    expected = {(c, s, r): args for c, s, r, args in mapping_cells}
    assert len(expected) >= 190, f"映射表三参 PREV 格只解析出 {len(expected)} 个，解析疑似失效"
    assert set(entries) == set(expected), (
        f"缺条目 {sorted(set(expected) - set(entries))}；多出（失效）条目 {sorted(set(entries) - set(expected))}")
    wrong = {k: (e.prev, expected[k]) for k, e in entries.items() if e.prev != expected[k]}
    assert not wrong, f"定位表记的取数参数与映射表不一致：{wrong}"


def test_targets_match_system_templates():
    """目标格的列头路径 / 行标签 / 公式标记与系统模板逐格一致；动作与公式标记一致；模板指纹一致。"""
    import json

    doc = json.loads(prev3.LOCATOR_FILE.read_text(encoding="utf-8"))
    books: dict[str, object] = {}
    bad: list[str] = []
    checked = 0
    for (code, sheet, ref), e in _entries().items():
        if e.action == "none":
            continue
        path = BACKEND / e.template
        assert hashlib.sha256(path.read_bytes()).hexdigest() == doc["templates"][e.template], f"{e.template} 指纹不符"
        if e.template not in books:
            books[e.template] = load_workbook(str(path))
        ws = next((w for w in books[e.template].worksheets if w.title.strip() == sheet), None)  # type: ignore[attr-defined]
        if ws is None:
            bad.append(f"{code}|{sheet}|{ref}: 模板里没有该表")
            continue
        header, label, is_formula = prev3.target_signature(ws, e.cell or "")
        if (header, label, is_formula) != (e.header, e.row_label, e.cell_is_formula):
            bad.append(f"{code}|{sheet}|{ref} {e.cell}: 模板 {header}/{label}/{is_formula} ≠ 定位表 "
                       f"{e.header}/{e.row_label}/{e.cell_is_formula}")
        if e.action != ("verify" if e.cell_is_formula else "write"):
            bad.append(f"{code}|{sheet}|{ref} {e.cell}: 动作 {e.action} 与公式标记 {e.cell_is_formula} 不一致")
        checked += 1
    assert checked >= 150, f"只核对了 {checked} 格，判据空转"
    assert not bad, "\n".join(bad)


def test_targets_have_prior_semantics(gen):
    """目标格属上年 / 期初语义：列头是非比较类的上年列，或行标签写明上年。"""
    assert not gen.has_prior_semantics(("期末与期初审定数比较", "变动额"), "合计"), "比较列必须判否（自检）"
    assert not gen.has_prior_semantics(("期末数", "审定数"), "合计"), "本期列必须判否（自检）"
    assert gen.has_prior_semantics(("期初数", "审定数"), "合计") and gen.has_prior_semantics(("本期数", "审定数"),
                                                                                                "一、上年年末余额")
    bad = [f"{k} {e.cell} 列头「{'/'.join(e.header)}」行「{e.row_label}」" for k, e in _entries().items()
           if e.action != "none" and not gen.has_prior_semantics(e.header, e.row_label)]
    assert not bad, f"目标格没有上年 / 期初语义：{bad}"


def test_none_entries_state_reason_and_manual_entries_are_live(gen, mapping_cells):
    silent = [k for k, e in _entries().items() if e.action == "none" and not e.reason.strip()]
    assert not silent, f"动作 none 的格必须写明原因：{silent}"
    live = {(c, s, r) for c, s, r, _ in mapping_cells}
    assert set(gen.MANUAL) <= live, f"人工裁定里有映射表已没有的格：{sorted(set(gen.MANUAL) - live)}"
    assert all((v.get("reason") or "").strip() for v in gen.MANUAL.values()), "人工裁定每条都必须写原因"


def test_shared_target_has_same_args():
    """同一目标格只许一套取数参数（否则只落第一格，另一格的值被静默吞掉）。"""
    seen: dict[tuple[str, str, str], tuple] = {}
    bad: list[str] = []
    for (code, sheet, ref), e in _entries().items():
        if e.action == "none":
            continue
        key = (code, sheet, e.cell or "")
        if key in seen and seen[key] != e.prev:
            bad.append(f"{key}: {seen[key]} vs {e.prev}（{ref}）")
        seen.setdefault(key, e.prev)
    assert not bad, bad


def test_template_header_change_breaks_check(gen, built):
    """模板列头改了 ⇒ 生成结果变化（``--check`` 红）：改 D1 审定表期初「审定数」列头后重扫。"""
    _path, ws = gen._find_sheet("D1", "审定表D1-1")
    original = ws["E6"].value
    assert original == "审定数", f"模板 D1 审定表 E6 是「{original}」，本变异的前提不成立"
    try:
        ws["E6"].value = "未审数"
        doc, unresolved = gen.build()
        assert unresolved or gen.render(doc) != prev3.LOCATOR_FILE.read_text(encoding="utf-8"), "改了列头生成结果不变"
    finally:
        ws["E6"].value = original


# ── 落格（模板副本） ─────────────────────────────────────────────────────────────────


def _template_copy(tmp_path: Path, wp_code: str, sheet: str) -> Path:
    entry = next(e for e in prev3.entries_for(wp_code) if e.sheet == sheet and e.action != "none")
    target = tmp_path / f"{wp_code}.xlsx"
    shutil.copy2(BACKEND / entry.template, target)
    return target


D4_AUDIT, D4_LISTED, D4_SOE = "营业收入审定表D4-1", "附注披露信息（上市公司）", "附注披露信息（国企）"


def _ws(wb, title: str):
    return next(w for w in wb.worksheets if w.title.strip() == title)


def test_apply_write_verify_missing_and_dedup(tmp_path):
    path = _template_copy(tmp_path, "D4", D4_AUDIT)
    wb = load_workbook(str(path))
    assert _ws(wb, D4_LISTED)["D10"].value is None and str(_ws(wb, D4_AUDIT)["I19"].value).startswith("=")
    written = prev3.apply_prefills(wb, "D4", {
        f"{D4_LISTED}!上年主营业务收入": prev3.PrevPrefill(Decimal("1234.5"), note="本项目 2024 年数据按「X」求值"),
        f"{D4_AUDIT}!上年审定数": prev3.PrevPrefill(Decimal("99"), note="本项目 2024 年数据按「Y」求值"),
        f"{D4_AUDIT}!期初余额": prev3.PrevPrefill(Decimal("99"), note="本项目 2024 年数据按「Y」求值"),
        f"{D4_SOE}!上年主营业务收入": prev3.PrevPrefill(None, "本项目没有 2024 年试算表，没有上年项目"),
    })
    assert written == 1
    listed, audit, soe = _ws(wb, D4_LISTED), _ws(wb, D4_AUDIT), _ws(wb, D4_SOE)
    assert listed["D10"].value == 1234.5 and listed["D10"].fill.start_color.rgb.endswith("E8F4FD")
    assert "1,234.50" in listed["D10"].comment.text and "上年同底稿取数" in listed["D10"].comment.text
    assert str(audit["I19"].value).startswith("="), "模板公式格不得改写"
    assert audit["I19"].comment.text.count("核对") == 1, "同一目标格（期初余额 / 上年审定数）只落一次"
    assert "99.00" in audit["I19"].comment.text and "未改写" in audit["I19"].comment.text
    assert soe["D9"].value is None, "取不到不写值（不写 0）"
    assert "没有上年项目" in soe["D9"].comment.text


def test_apply_skips_mismatched_template_and_user_cells(tmp_path):
    path = _template_copy(tmp_path, "D4", D4_AUDIT)
    wb = load_workbook(str(path))
    audit, listed = _ws(wb, D4_AUDIT), _ws(wb, D4_LISTED)
    audit["I6"].value = "未审数"  # 模拟事务所覆盖版模板改了列头：落格前核对不符 ⇒ 不动
    prefills = {f"{D4_AUDIT}!上年审定数": prev3.PrevPrefill(Decimal("1")),
                f"{D4_LISTED}!上年主营业务收入": prev3.PrevPrefill(Decimal("2"))}
    written = prev3.apply_prefills(wb, "D4", prefills, skip_cells={(D4_LISTED, "D10")})
    assert written == 0
    assert audit["I19"].comment is None, "列头与定位表不符必须不落格"
    assert listed["D10"].value is None and listed["D10"].comment is None, "用户自定义公式格不动"


def test_prefill_xlsx_no_keyword_guess_for_prev3(tmp_path):
    """初次预填：三参格不再按关键词猜格（旧规则把「上年审定数」写进本期审定数列 E7）；改成非 PREV 公式的
    三参格按改后的公式取数、落定位表的格。"""
    from app.services.wp_template_init_service import prefill_workpaper_xlsx

    path = _template_copy(tmp_path, "D4", D4_AUDIT)
    tb = {"6001": {"期末余额": 500.0, "期初余额": 70.0, "本期发生额": 500.0}}
    prefill_workpaper_xlsx(path, "D4", tb, mapping_overrides={f"{D4_AUDIT}!上年审定数": "=TB('6001','期初余额')"},
                           prev_values={f"{D4_LISTED}!上年主营业务收入": prev3.PrevPrefill(Decimal("8"))})
    wb = load_workbook(str(path))
    audit, listed = _ws(wb, D4_AUDIT), _ws(wb, D4_LISTED)
    assert audit["E7"].value is None, "上年数被按关键词写进本期审定数列"
    assert str(audit["I19"].value).startswith("=") and "70.00" in audit["I19"].comment.text
    assert "改写后的公式取数核对" in audit["I19"].comment.text
    assert listed["D10"].value == 8


def test_prefill_xlsx_without_prev_values_leaves_prev3_cells(tmp_path):
    from app.services.wp_template_init_service import prefill_workpaper_xlsx

    path = _template_copy(tmp_path, "D4", D4_AUDIT)
    prefill_workpaper_xlsx(path, "D4", {"6001": {"期末余额": 1.0}})
    audit = _ws(load_workbook(str(path)), D4_AUDIT)
    assert audit["E7"].value is None and audit["I19"].comment is None


def test_locate_cell_uses_locator_only_for_prev3(tmp_path, monkeypatch):
    """重新预填落格：三参格只认定位表（write 给坐标、verify / 定位表缺失不落格），其余语义格行为不变。"""
    from app.services.formula_settings.wp_effective import locate_cell

    k6 = load_workbook(str(_template_copy(tmp_path, "K6", "附注披露信息（上市公司）")))
    d1 = load_workbook(str(_template_copy(tmp_path, "D1", "审定表D1-1")))
    k6_notes, d1_audit = _ws(k6, "附注披露信息（上市公司）"), _ws(d1, "审定表D1-1")
    assert locate_cell(k6_notes, "上年账面价值_上市", "K6") == "G16"
    assert locate_cell(d1_audit, "上年审定数", "D1") is None, "模板公式格不落格"
    assert locate_cell(d1_audit, "期初余额", "D1") == "B7", "非三参格仍按原规则（本 spec 不改）"
    monkeypatch.setattr(prev3, "load_locator", lambda: {})
    assert locate_cell(k6_notes, "上年账面价值_上市", "K6") is None, "定位表缺失时不得退回关键词猜格"


# ── 取值与路由（SQLite 真 ORM） ─────────────────────────────────────────────────────────


def _extra_models():
    from app.models.audit_platform_models import TrialBalance
    from app.models.dataset_models import LedgerDataset
    from app.models.workpaper_models import WorkingPaper, WpFormula, WpIndex

    return (TrialBalance, LedgerDataset, WpIndex, WorkingPaper, WpFormula)


async def _world(db, *, cur_year=2025, prior_tb=True, cur_tb=True):
    """本年项目 + 同企业代码的上年项目（2024 试算表 1121：审定 800、期初 100）+ 本年 D1 底稿。"""
    from app.models.audit_platform_models import AccountCategory, TrialBalance
    from app.models.base import ProjectStatus
    from app.models.core import Project
    from app.models.workpaper_models import WorkingPaper, WpIndex

    def project(name, year):
        pid = uuid.uuid4()
        db.add(Project(id=pid, name=name, client_name=name, status=ProjectStatus.execution, template_type="soe",
                       report_scope="standalone", audit_year=year, company_code="C-PREV3"))
        return pid

    def tb(pid, year, code, audited, opening="0", deleted=False):
        return TrialBalance(project_id=pid, year=year, company_code="001", standard_account_code=code,
                            account_name=code, account_category=AccountCategory.asset, is_deleted=deleted,
                            unadjusted_amount=Decimal(audited), audited_amount=Decimal(audited),
                            opening_balance=Decimal(opening))

    cur, prior = project("【测试】本年", cur_year), project("【测试】上年", 2024)
    await db.flush()
    if cur_tb:
        db.add(tb(cur, 2025, "1121", "900", opening="800"))
    if prior_tb:
        db.add(tb(prior, 2024, "1121", "800", opening="100"))
    idx = WpIndex(id=uuid.uuid4(), project_id=cur, wp_code="D1", wp_name="应收票据")
    db.add(idx)
    await db.flush()
    wp = WorkingPaper(id=uuid.uuid4(), project_id=cur, wp_index_id=idx.id, file_path="x.xlsx", source_type="template")
    db.add(wp)
    await db.commit()
    return cur, wp.id


async def test_resolve_prev_prefill_values_and_reasons():
    from tests.formula_settings._env import make_db

    key = "审定表D1-1!上年审定数"
    async with make_db(extra_models=_extra_models()) as factory, factory() as db:
        cur, _wp = await _world(db)
        got = await prev3.resolve_prev_prefill(db, cur, "D1")
        assert got[key].value == Decimal("800") and "上年项目「【测试】上年」 2024 年数据" in got[key].note
        assert "TB('1121','审定数')" in got[key].note
        changed = await prev3.resolve_prev_prefill(db, cur, "D1", {key: "=PREV('D1','审定表D1-1','期初余额')"})
        assert changed[key].value == Decimal("100"), "改成别的三参 PREV 时按改后的参数取"
        assert key not in await prev3.resolve_prev_prefill(db, cur, "D1", {key: "=TB('1121','期初余额')"})
    async with make_db(extra_models=_extra_models()) as factory, factory() as db:
        cur, _wp = await _world(db, prior_tb=False)
        got = await prev3.resolve_prev_prefill(db, cur, "D1")
        assert got[key].value is None and "没有 2024 年试算表" in got[key].reason, got[key]
    async with make_db(extra_models=_extra_models()) as factory, factory() as db:
        cur, _wp = await _world(db, cur_year=None)
        got = await prev3.resolve_prev_prefill(db, cur, "D1")
        assert got[key] == prev3.PrevPrefill(None, "项目没有审计年度，无法确定上年")


async def test_tb_data_for_prefill_filters_year_and_deleted(monkeypatch):
    """初次预填取本年试算表：只取审计年度、排除已删除行（原 SQL 不按年度也不排除删除）。"""
    from app.models.audit_platform_models import AccountCategory, TrialBalance
    from app.routers import wp_template_files
    from tests.formula_settings._env import make_db

    async def no_occurrence(*_a, **_k):
        return {}

    monkeypatch.setattr(wp_template_files, "fetch_occurrence_by_standard_code", no_occurrence)
    async with make_db(extra_models=_extra_models()) as factory, factory() as db:
        cur, _wp = await _world(db)
        for year, code, amount, deleted in ((2024, "1001", "99", False), (2025, "1002", "5", True)):
            db.add(TrialBalance(project_id=cur, year=year, company_code="001", standard_account_code=code,
                                account_name=code, account_category=AccountCategory.asset, is_deleted=deleted,
                                unadjusted_amount=Decimal(amount), audited_amount=Decimal(amount)))
        await db.commit()
        data = await wp_template_files._get_tb_data_for_prefill(db, cur)
        assert set(data) == {"1121"}, f"混入了其他年度 / 已删除行：{sorted(data)}"
        assert data["1121"]["审定数"] == 900.0 and data["1121"]["期初余额"] == 800.0


async def test_router_first_init_annotates_prior_year(tmp_path, monkeypatch):
    """路由端到端：首次打开底稿 → 从系统模板初始化 → 三参格按定位表批注上年值（来源上年项目），不写进猜出的列。"""
    from app.routers import wp_template_files
    from app.services import wp_template_init_service
    from tests.formula_settings._env import make_db

    async def no_occurrence(*_a, **_k):
        return {}

    storage = tmp_path / "wp.xlsx"
    monkeypatch.setattr(wp_template_files, "fetch_occurrence_by_standard_code", no_occurrence)
    monkeypatch.setattr(wp_template_files, "get_workpaper_file", lambda *_a: None)
    monkeypatch.setattr(wp_template_init_service, "get_workpaper_storage_path", lambda *_a: storage)
    async with make_db(extra_models=_extra_models()) as factory, factory() as db:
        cur, wp_id = await _world(db)
        response = await wp_template_files.get_workpaper_xlsx(project_id=cur.hex, wp_id=wp_id.hex, db=db)
    assert Path(response.path) == storage
    audit = _ws(load_workbook(str(storage)), "审定表D1-1")
    assert str(audit["E10"].value).startswith("=") and "800.00" in audit["E10"].comment.text
    assert "上年项目「【测试】上年」" in audit["E10"].comment.text
    assert audit["J7"].value is None, "旧规则把上年审定数写进「期末与期初审定数比较」列"

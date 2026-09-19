"""D4-9/10/11 导入导出：英文字段映射 + D4-10 公式元数据不被 I/O 冲掉。"""
from __future__ import annotations

import io
import json

from openpyxl import Workbook

from app.routers.wp_render_strategies import _d4_import_export as m


def test_parse_d4_10_maps_chinese_headers_to_english_keys():
    headers = m._SHEET_HEADERS["D4-10"]
    row = ("客户甲", "产品A", 1000, 10, 100, 90, "季节性", 95, "市价波动")
    parsed = m._parse_d4_10_row(row, headers)
    assert parsed["customer"] == "客户甲"
    assert parsed["product"] == "产品A"
    assert parsed["amount"] == 1000
    assert parsed["quantity"] == 10
    assert parsed["unitPrice"] == 100
    assert parsed["avgPrice"] == 90
    assert parsed["avgReason"] == "季节性"
    assert parsed["marketPrice"] == 95
    assert parsed["marketReason"] == "市价波动"


def test_parse_d4_9_and_d4_11_field_mapping():
    h9 = m._SHEET_HEADERS["D4-9"]
    p9 = m._parse_d4_9_row(("客户乙", 200, 5, "3"), h9)
    assert p9 == {"name": "客户乙", "amount": 200.0, "quantity": 5.0, "priorRank": "3"}

    h11 = m._SHEET_HEADERS["D4-11"]
    p11 = m._parse_d4_11_row(
        ("客户丙", "规格X", 12, 3, "2024-01-01", "SO-1", "2024-01-02", 10, 11, "议价", "IDX-1", "备注"),
        h11,
    )
    assert p11["customer"] == "客户丙"
    assert p11["product"] == "规格X"
    assert p11["unitPrice"] == 12
    assert p11["listPrice"] == 10
    assert p11["priceSource"] == "IDX-1"


def test_merge_d4_10_import_rows_preserves_formula_metadata():
    """真实 merge helper：只换 rows，公式覆盖位与总额保留。"""
    existing = {
        "rows": [{"customer": "旧", "product": "P", "amount": 1}],
        "totalAmount": 99999,
        "totalQuantity": 42,
        "totalAmountManualOverride": False,
        "totalAmountFormulaRef": "WP('D4-2','本期未审合计')",
    }
    imported_rows = [{"customer": "新", "product": "Q", "amount": 5}]
    merged = m.merge_d4_10_import_rows(existing, imported_rows)
    assert merged["rows"] == imported_rows
    assert merged["totalAmount"] == 99999
    assert merged["totalAmountManualOverride"] is False
    assert "WP('D4-2'" in merged["totalAmountFormulaRef"]
    roundtrip = json.loads(json.dumps(merged, ensure_ascii=False))
    assert isinstance(roundtrip, dict)
    assert "rows" in roundtrip


def test_merge_d4_10_prefers_xlsx_formula_meta_sheet():
    """导出副表「公式元数据」在导入时优先于既有 store 的空缺位。"""
    existing = {"rows": [], "totalAmount": 1}
    imported = [{"customer": "C", "product": "P", "amount": 2}]
    meta = {
        "totalAmount": 555,
        "totalQuantity": 7,
        "totalAmountManualOverride": True,
        "totalAmountFormulaRef": "WP('D4-2','本期未审合计')",
    }
    merged = m.merge_d4_10_import_rows(existing, imported, meta)
    assert merged["totalAmount"] == 555
    assert merged["totalQuantity"] == 7
    assert merged["totalAmountManualOverride"] is True
    assert merged["totalAmountFormulaRef"] == "WP('D4-2','本期未审合计')"


def test_extract_d4_10_formula_meta_from_exported_workbook():
    """往返：写入「公式元数据」副表后能抽回。"""
    wb = Workbook()
    ws = wb.active
    ws.title = "D4-10"
    ws.append(m._SHEET_HEADERS["D4-10"])
    meta_ws = wb.create_sheet("公式元数据")
    meta_ws.append(["字段", "值"])
    meta_ws.append(["totalAmount", 12345])
    meta_ws.append(["totalQuantity", 9])
    meta_ws.append(["totalAmountManualOverride", False])
    meta_ws.append(["totalAmountFormulaRef", "WP('D4-2','本期未审合计')"])
    buf = io.BytesIO()
    wb.save(buf)
    buf.seek(0)

    from openpyxl import load_workbook

    loaded = load_workbook(buf, read_only=True, data_only=True)
    meta = m._extract_d4_10_formula_meta_from_wb(loaded)
    loaded.close()
    assert meta["totalAmount"] == 12345
    assert meta["totalQuantity"] == 9
    assert "WP('D4-2'" in str(meta["totalAmountFormulaRef"])

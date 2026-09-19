"""交付导出契约单元测试（Task 10.1 / Req 18.1-18.5）。

覆盖 `delivery_export`：
- 公式解析为静态值（SUM / 四则运算 / 跨公式引用），产物不留 '=' 表达式。
- 悬空引用（未知函数）降级导出并记入 dangling 清单。
- RFC 5987 中文文件名编码可解码还原（往返）。
"""

from __future__ import annotations

from urllib.parse import unquote

from openpyxl import Workbook

from app.services.formula_management.delivery_export import (
    content_disposition_attachment,
    flatten_workbook_formulas,
)


def _build_wb() -> Workbook:
    wb = Workbook()
    ws = wb.active
    ws.title = "资产负债表"
    ws["C5"] = 100.0
    ws["C6"] = 200.0
    ws["C7"] = -50.0
    ws["C8"] = "=SUM(C5:C7)"      # 250
    ws["C9"] = "=C8*0.25"          # 62.5（引用另一公式格）
    ws["C10"] = "=C5+C6-C7"        # 350
    ws["C11"] = "=SUM(C5:C7,C9)"   # 312.5（区间 + 单格混合）
    return wb


def test_flatten_resolves_formulas_to_static_values():
    wb = _build_wb()
    result = flatten_workbook_formulas(wb)
    ws = wb["资产负债表"]

    assert ws["C8"].value == 250.0
    assert ws["C9"].value == 62.5
    assert ws["C10"].value == 350.0
    assert ws["C11"].value == 312.5
    assert result.flattened == 4
    assert result.dangling_count == 0


def test_flatten_leaves_no_residual_formula_expression():
    """Req 18.3：交付产物不含以 '=' 开头的可重算表达式。"""
    wb = _build_wb()
    flatten_workbook_formulas(wb)
    for ws in wb.worksheets:
        for row in ws.iter_rows():
            for cell in row:
                assert not (
                    isinstance(cell.value, str) and cell.value.startswith("=")
                ), f"残留公式：{ws.title}!{cell.coordinate} = {cell.value!r}"


def test_flatten_dangling_reference_downgraded_and_logged():
    """Req 18.4：无法求值的公式降级（清空/最近计算值）并记入 dangling 清单。"""
    wb = Workbook()
    ws = wb.active
    ws.title = "利润表"
    ws["B5"] = 10.0
    ws["B6"] = "=BADFUNC(B5)"  # 未知函数 → 悬空

    result = flatten_workbook_formulas(wb)

    # 悬空单元不再是公式表达式
    assert not (isinstance(ws["B6"].value, str) and ws["B6"].value.startswith("="))
    assert result.dangling_count == 1
    dangling = result.dangling[0]
    assert dangling.location == "利润表!B6"
    assert dangling.formula == "=BADFUNC(B5)"


def test_flatten_text_and_blank_cells_treated_as_zero_in_arithmetic():
    wb = Workbook()
    ws = wb.active
    ws["A1"] = "标签文字"
    ws["A2"] = 5.0
    ws["A3"] = "=A1+A2"  # 文字视作 0 → 5
    ws["A4"] = "=SUM(A1:A2)"  # 文字忽略 → 5
    flatten_workbook_formulas(wb)
    assert ws["A3"].value == 5.0
    assert ws["A4"].value == 5.0


def test_content_disposition_rfc5987_roundtrip():
    """Req 18.5：中文文件名 RFC 5987 编码后可解码还原。"""
    filename = "某某公司_2024年度财务报表(审定).xlsx"
    header = content_disposition_attachment(filename, ascii_fallback="report.xlsx")

    assert header.startswith("attachment; ")
    assert 'filename="report.xlsx"' in header
    assert "filename*=UTF-8''" in header

    encoded = header.split("filename*=UTF-8''", 1)[1]
    assert unquote(encoded) == filename


def test_content_disposition_default_ascii_fallback():
    header = content_disposition_attachment("附注.docx")
    # 纯中文名剥离 ASCII 后回退到默认名，仍保留 filename* 真实名
    assert 'filename="download"' in header or 'filename=".docx"' in header
    encoded = header.split("filename*=UTF-8''", 1)[1]
    assert unquote(encoded) == "附注.docx"

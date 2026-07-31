"""N1 公式预设纯净性守卫（Property 12）。

spec: `.kiro/specs/n1-four-table-extraction-and-disclosure-alignment/` R6

背景（修掉的真实缺陷）
----------------------
`prefill_formula_mapping.json` 曾有一个 `wp_code=N1` 却名为「税金分析程序」的块：
- `sheet` 写 `分析程序N1-3`，而 N1 源模板该位置是 **调整分录汇总N1-3**（无「分析程序」sheet）；
- `account_codes` 是 `2221`（应交税费）/ `6401` / `6403`（税金及附加）—— 与递延所得税资产无关，
  且 N4 已有自己的税金及附加块；
- 表达式 `TB_SUM('2221~6403','期末余额')` 是**跨科目大类的病态区间**
  （从负债 2221 一路到损益 6403，把权益/成本/损益全括进去）；
- 它的 `cell_ref` 「上年审定数」与 N1-1 块同名 —— 而 `preset_library.convert_prefill_presets`
  的 `page_key = f"workpaper:{wp_code}"` **忽略 sheet** → 同一页两条同名预设互相遮蔽。

因此 N1 底稿页的公式管理里会出现应交税费预设，且「上年审定数」指向不确定。
"""

from __future__ import annotations

import json
import re
from pathlib import Path

import pytest

BACKEND = Path(__file__).resolve().parents[2]
MAPPING_PATH = BACKEND / "data" / "prefill_formula_mapping.json"

# N1 允许引用的科目（每一条都有源模板依据）：
#   1811.*  递延所得税资产及其子科目（N1-1 审定表 / N1-2 明细 / N1-4 测算表第 I 列）
#   2901    递延所得税负债（N1-4 测算表第 M 列「递延所得税负债期末账面余额（5）」）
#   4104    未分配利润（N1-5 R14「期末未分配利润」的账面金额列；报表行 BS-088 同口径）
ALLOWED_PREFIXES = ("1811", "2901", "4104")

# 明确禁止的科目（曾被误挂进 N1，属 N2/N4 内容）
FORBIDDEN_CODES = {"2221", "6401", "6403"}

_TB_RE = re.compile(r"TB\w*\('([^']+)'")


@pytest.fixture(scope="module")
def n1_blocks() -> list[dict]:
    data = json.loads(MAPPING_PATH.read_text(encoding="utf-8"))
    blocks = [m for m in data.get("mappings", []) if m.get("wp_code") == "N1"]
    assert blocks, "prefill_formula_mapping.json 中应存在 wp_code=N1 的块"
    return blocks


def _cells(blocks: list[dict]) -> list[tuple[str, str, str]]:
    """→ [(sheet, cell_ref, formula)]，跳过 formula 为 null 的占位项。"""
    out: list[tuple[str, str, str]] = []
    for b in blocks:
        for c in b.get("cells") or []:
            f = c.get("formula")
            if not f:
                continue
            out.append((str(b.get("sheet") or ""), str(c.get("cell_ref") or ""), str(f)))
    return out


def _codes(formula: str) -> list[str]:
    return [m.strip() for m in _TB_RE.findall(formula)]


# ─── 反向自检：先证明提取逻辑非空转 ──────────────────────────────────────────


def test_self_check_extraction_is_non_empty(n1_blocks):
    cells = _cells(n1_blocks)
    assert len(cells) >= 20, f"N1 预设 cell 数异常偏少（{len(cells)}），提取逻辑可能失效"
    all_codes = {c for _, _, f in cells for c in _codes(f)}
    assert "1811" in all_codes, "未提取到科目 1811 → 正则失效"


# ─── Property 12: 科目纯净 ───────────────────────────────────────────────────


def test_no_forbidden_accounts_in_formulas(n1_blocks):
    bad: list[str] = []
    for sheet, ref, f in _cells(n1_blocks):
        for code in _codes(f):
            head = code.split("~")[0]
            if head in FORBIDDEN_CODES or code.split("~")[-1] in FORBIDDEN_CODES:
                bad.append(f"{sheet}::{ref}::{f}")
    assert bad == [], f"N1 预设引用了应交税费/税金及附加科目（属 N2/N4）：{bad}"


def test_no_forbidden_accounts_in_block_metadata(n1_blocks):
    bad = [
        f"{b.get('sheet')}::{c}"
        for b in n1_blocks
        for c in (b.get("account_codes") or [])
        if str(c) in FORBIDDEN_CODES
    ]
    assert bad == [], f"N1 块的 account_codes 含禁用科目：{bad}"


def test_all_referenced_accounts_are_whitelisted(n1_blocks):
    bad: list[str] = []
    for sheet, ref, f in _cells(n1_blocks):
        for code in _codes(f):
            for part in code.split("~"):
                if not part.startswith(ALLOWED_PREFIXES):
                    bad.append(f"{sheet}::{ref}::{part}")
    assert bad == [], (
        f"N1 预设引用了白名单外科目：{bad}；"
        "新增科目必须先在源模板中找到依据并加入 ALLOWED_PREFIXES 并写明理由"
    )


def test_no_cross_category_range_formulas(n1_blocks):
    """区间求和的首尾必须属同一科目大类（首位数字相同），防 `2221~6403` 式病态区间。"""
    bad: list[str] = []
    for sheet, ref, f in _cells(n1_blocks):
        for code in _codes(f):
            if "~" not in code:
                continue
            lo, hi = (p.strip() for p in code.split("~", 1))
            if not lo or not hi or lo[0] != hi[0]:
                bad.append(f"{sheet}::{ref}::{code}")
    assert bad == [], f"跨科目大类的区间求和（会把整段科目表括进来）：{bad}"


# ─── cell_ref 唯一性（page_key 忽略 sheet → 同名必遮蔽）─────────────────────


def test_cell_refs_unique_within_n1(n1_blocks):
    """`convert_prefill_presets` 的 page_key = `workpaper:N1`，**sheet 不参与** →
    同名 cell_ref 在公式管理里互相遮蔽，用户看到哪条不确定。"""
    seen: dict[str, list[str]] = {}
    for sheet, ref, f in _cells(n1_blocks):
        seen.setdefault(ref, []).append(f"{sheet}={f}")
    dups = {k: v for k, v in seen.items() if len(v) > 1}
    assert dups == {}, f"N1 内 cell_ref 撞键（page_key 忽略 sheet）：{dups}"


# ─── sheet 名必须是 N1 真实 sheet ────────────────────────────────────────────


def test_sheets_exist_in_source_template(n1_blocks):
    """块的 `sheet` 必须能对上 N1 源模板的 sheet（允许省略科目前缀的简写）。"""
    from openpyxl import load_workbook

    xlsx = BACKEND / "wp_templates" / "N" / "N1 递延所得税资产.xlsx"
    assert xlsx.exists(), f"源模板缺失：{xlsx}"
    real = load_workbook(xlsx, read_only=True).sheetnames
    assert len(real) >= 9, f"源模板 sheet 数异常（{len(real)}）→ 自检失败"

    bad: list[str] = []
    for b in n1_blocks:
        s = str(b.get("sheet") or "").strip()
        if not s:
            bad.append("(空 sheet)")
            continue
        # 简写容忍：真实 tab 名以该串结尾即可（如「审定表N1-1」↔「递延所得税资产审定表N1-1」）
        if not any(r == s or r.endswith(s) for r in real):
            bad.append(s)
    assert bad == [], (
        f"N1 预设块的 sheet 在源模板中不存在：{bad}；实际 sheet={real}"
    )


# ─── 表达式可解析 ────────────────────────────────────────────────────────────


# `prefill_formula_mapping` 用的是 **prefill 引擎**的词汇表，比 report 公式引擎多几个函数。
# `ADJ()` 由 `prefill_engine._resolve_adj_formula` 实现，未注册进 `formula_engine._REGISTRY`
# → `validate_formula` 会把它报成「未知函数」。这是既有形态（N1-1 的 AJE/RJE 调整预设一直如此），
# 不是本 spec 引入的，故按 prefill 词汇表放行。
_PREFILL_ONLY_FUNCS = {"ADJ"}
_UNKNOWN_FUNC_RE = re.compile(r"未知函数:\s*([A-Z_]+)\(\)")


def test_all_formulas_parse(n1_blocks):
    from app.services.formula_engine import validate_formula

    bad: list[str] = []
    for sheet, ref, f in _cells(n1_blocks):
        expr = f[1:] if f.startswith("=") else f
        try:
            # `validate_formula` 返回**错误列表**（空 = 合法）
            errors = validate_formula(expr)
        except Exception as e:  # noqa: BLE001
            bad.append(f"{sheet}::{ref}::{f} → 异常 {e}")
            continue
        real = []
        for msg in errors:
            m = _UNKNOWN_FUNC_RE.search(msg)
            if m and m.group(1) in _PREFILL_ONLY_FUNCS:
                continue  # prefill 专属函数，非缺陷
            real.append(msg)
        if real:
            bad.append(f"{sheet}::{ref}::{f} → {real}")
    assert bad == [], f"N1 预设存在无法解析的表达式：{bad}"


def test_prefill_only_func_allowlist_is_still_needed():
    """反向自检：若 ADJ 某天注册进 formula_engine，本豁免应被移除（防豁免长期挂着）。"""
    from app.services.formula_engine import _REGISTRY

    known = set(_REGISTRY.known_function_names())
    stale = _PREFILL_ONLY_FUNCS & known
    assert stale == set(), (
        f"{stale} 已注册进 formula_engine → 请从 _PREFILL_ONLY_FUNCS 移除该豁免"
    )


def test_validate_formula_guard_is_not_vacuous():
    """反向自检：断言校验器确实能报错，否则上一条测试恒绿。"""
    from app.services.formula_engine import validate_formula

    assert validate_formula("TB('1811','期末余额'") != []  # 缺右括号
    assert validate_formula("NOPE('1811')") != []  # 未知函数
    assert validate_formula("TB('1811','期末余额')") == []


# ─── 经 convert_prefill_presets 落地后的最终形态 ─────────────────────────────


def test_converted_presets_are_clean():
    from app.services.formula_management.preset_library import convert_prefill_presets

    entries = [e for e in convert_prefill_presets() if e.page_key == "workpaper:N1"]
    assert entries, "workpaper:N1 应有预设条目"
    for e in entries:
        for code in _codes(e.expression):
            assert code.split("~")[0].startswith(ALLOWED_PREFIXES), (
                f"公式管理里 N1 出现白名单外科目：{e.target_cell} → {e.expression}"
            )
    refs = [e.target_cell for e in entries]
    assert len(refs) == len(set(refs)), f"workpaper:N1 预设 target_cell 撞键：{refs}"


def test_n1_4_and_n1_5_have_presets(n1_blocks):
    """N1-4 测算表与 N1-5 亏损检查表必须有取数预设（原先完全没有）。"""
    sheets = {str(b.get("sheet") or "") for b in n1_blocks}
    assert any("N1-4" in s for s in sheets), "N1-4 测算表缺公式预设"
    assert any("N1-5" in s for s in sheets), "N1-5 亏损检查表缺公式预设"


def test_n1_2_has_both_opening_and_closing(n1_blocks):
    """N1-2 的 1811.01~.07 期初 / 期末余额预设必须成对（源模板期初段与期末段同构）。"""
    refs = {ref for sheet, ref, _ in _cells(n1_blocks) if "N1-2" in sheet}
    for name in ("公允价值变动", "资产减值准备", "长期职工薪酬", "可抵扣亏损",
                 "预提费用", "递延收益", "租赁年金"):
        assert f"{name}_期末余额" in refs, f"N1-2 缺 {name}_期末余额"
        assert f"{name}_期初余额" in refs, f"N1-2 缺 {name}_期初余额"

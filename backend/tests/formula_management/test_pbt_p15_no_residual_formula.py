"""属性测试 P15：交付导出无残留公式表达式（Task 10.2 / Req 18.1, 18.2, 18.3）。

被测：`app.services.formula_management.delivery_export.flatten_workbook_formulas`
（Task 10.1 已实现）。用 Hypothesis 随机构造含公式的 openpyxl 工作簿（SUM / 单元
引用 / 四则运算 / 跨公式引用），flatten 后断言产物单元不含任何以 `=` 开头的可重算
公式表达式，全部解析为静态数值。
"""

from __future__ import annotations

from hypothesis import given, settings
from hypothesis import strategies as st
from openpyxl import Workbook
from openpyxl.utils import get_column_letter

from app.services.formula_management.delivery_export import flatten_workbook_formulas

# 有限、量级受控的数值（避免 NaN/Inf 与乘法溢出，保证公式可解析为有限值）
_finite = st.floats(
    min_value=-1_000.0,
    max_value=1_000.0,
    allow_nan=False,
    allow_infinity=False,
)
# 公式内出现的非负数字字面量（正负号交给运算符，避免生成 "A1--3" 之类歧义 token）
_literal = st.one_of(
    st.integers(min_value=0, max_value=100).map(str),
    st.floats(
        min_value=0.0, max_value=100.0, allow_nan=False, allow_infinity=False
    ).map(lambda v: repr(round(v, 3))),
)
_BINOPS = ("+", "-", "*")  # 不含除法：避免除零导致悬空，确保全部可解析为静态值


def _build_expr(draw, coords: list[str], depth: int) -> str:
    """递归构造受限公式表达式（引用 / 数字字面量 / 四则 / 括号）。

    仅引用 ``coords`` 中已存在且已定义的单元，保证求值不悬空。
    """
    if depth <= 0 or draw(st.booleans()):
        # 叶子：单元引用或数字字面量
        return draw(st.one_of(st.sampled_from(coords), _literal))
    left = _build_expr(draw, coords, depth - 1)
    right = _build_expr(draw, coords, depth - 1)
    op = draw(st.sampled_from(_BINOPS))
    expr = f"{left}{op}{right}"
    if draw(st.booleans()):
        expr = f"({expr})"
    return expr


@st.composite
def _workbook_spec(draw):
    """生成 (数据值列表, 公式串列表)。

    - 数据放在 A 列（A1..An），承载随机静态数值。
    - 公式放在 B 列（B1..Bm），可引用 A 列任意单元及 **前序** B 列公式单元
      （B1..B(k-1)），从而覆盖跨公式引用且天然无环。
    """
    n_data = draw(st.integers(min_value=2, max_value=8))
    data = draw(st.lists(_finite, min_size=n_data, max_size=n_data))

    n_formula = draw(st.integers(min_value=1, max_value=6))
    a_coords = [f"A{r}" for r in range(1, n_data + 1)]
    formulas: list[str] = []
    for k in range(n_formula):
        # 可引用坐标 = 全部 A 列 + 已生成的前序 B 列公式单元
        coords = a_coords + [f"B{j}" for j in range(1, k + 1)]
        kind = draw(st.integers(min_value=0, max_value=2))
        if kind == 0:
            # SUM 连续区间
            a = draw(st.integers(min_value=1, max_value=n_data))
            b = draw(st.integers(min_value=1, max_value=n_data))
            form = f"=SUM(A{min(a, b)}:A{max(a, b)})"
        elif kind == 1:
            # SUM 逗号分隔引用/数字混合
            items = draw(
                st.lists(
                    st.one_of(st.sampled_from(coords), _literal),
                    min_size=1,
                    max_size=4,
                )
            )
            form = f"=SUM({','.join(items)})"
        else:
            # 一般四则运算表达式（引用 + 数字 + +-* + 括号）
            form = "=" + _build_expr(draw, coords, depth=3)
        formulas.append(form)
    return data, formulas


def _assemble_workbook(data: list[float], formulas: list[str]) -> Workbook:
    wb = Workbook()
    ws = wb.active
    ws.title = "资产负债表"
    for i, val in enumerate(data, start=1):
        ws.cell(row=i, column=1, value=val)  # A 列静态数据
    for k, form in enumerate(formulas, start=1):
        ws.cell(row=k, column=2, value=form)  # B 列公式
    return wb


# Feature: formula-management-library, Property 15: 交付导出无残留公式表达式 —
# 对任意含公式的报表或文档，交付导出后的产物不包含任何可被下游重新求值的公式表达式，
# 全部为解析后的静态值。
# Validates: Requirements 18.1, 18.2, 18.3
@settings(max_examples=100)
@given(_workbook_spec())
def test_delivery_export_leaves_no_residual_formula(spec):
    data, formulas = spec
    wb = _assemble_workbook(data, formulas)

    result = flatten_workbook_formulas(wb)

    # 18.3：交付产物任何单元都不再承载以 '=' 开头的可重算公式表达式。
    for ws in wb.worksheets:
        for row in ws.iter_rows():
            for cell in row:
                assert not (
                    isinstance(cell.value, str) and cell.value.startswith("=")
                ), f"残留公式：{ws.title}!{cell.coordinate} = {cell.value!r}"

    # 18.1/18.2：本生成器构造的公式均可解析 —— 全部被 flatten 为有限静态数值，
    # 无悬空降级。
    ws = wb["资产负债表"]
    assert result.dangling_count == 0
    assert result.flattened == len(formulas)
    for k in range(1, len(formulas) + 1):
        value = ws.cell(row=k, column=2).value
        assert isinstance(value, (int, float)) and not isinstance(value, bool), (
            f"公式单元 {get_column_letter(2)}{k} 未解析为数值：{value!r}"
        )

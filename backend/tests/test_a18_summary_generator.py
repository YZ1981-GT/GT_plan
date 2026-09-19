"""Tests for a18_summary_generator — A18-1 audit summary from A17 chapters."""

from __future__ import annotations

import uuid
from unittest.mock import AsyncMock, MagicMock

import pytest

from app.services.a18_summary_generator import generate_audit_summary


@pytest.mark.asyncio
async def test_generate_no_a17_wp():
    db = AsyncMock()
    db.execute = AsyncMock(return_value=MagicMock(scalar_one_or_none=MagicMock(return_value=None)))
    project_id = uuid.uuid4()

    result = await generate_audit_summary(db, project_id)

    assert result["completeness"] == 0.0
    assert result["summary_sections"] == []
    assert "未找到 A17-1" in (result.get("message") or "")


@pytest.mark.asyncio
async def test_generate_with_chapters_and_kam(monkeypatch):
    project_id = uuid.uuid4()
    a17_wp_id = uuid.uuid4()
    kam_wp_id = uuid.uuid4()

    async def fake_find(db, pid, code):
        if code == "A17-1":
            return a17_wp_id
        if code == "A17-2-1":
            return kam_wp_id
        return None

    monkeypatch.setattr("app.services.a18_summary_generator._find_wp_id", fake_find)

    async def fake_chapters(db, wp_id, ids):
        return {
            "A17-1-ch01": "约定范围内容",
            "A17-1-ch14": "无保留意见",
        }

    monkeypatch.setattr(
        "app.services.a18_summary_generator._load_chapter_remarks",
        fake_chapters,
    )

    async def fake_kam(db, pid):
        return "- 收入确认\n  确定为KAM的原因：重大判断"

    monkeypatch.setattr(
        "app.services.a18_summary_generator._load_kam_summary",
        fake_kam,
    )

    db = AsyncMock()
    result = await generate_audit_summary(db, project_id)

    assert result["completeness"] > 0
    assert any(s["content"] for s in result["summary_sections"])
    assert "收入确认" in result["formatted_text"]
    assert "关键审计事项" in result["formatted_text"]


def test_append_summary_to_document():
    from docx import Document
    from app.services.a18_summary_generator import append_summary_to_document

    doc = Document()
    doc.add_paragraph("抬头")
    doc.add_paragraph("致同会计师事务所（特殊普通合伙）")

    summary = {
        "summary_sections": [
            {"title": "审计结论", "content": "无保留意见"},
        ],
    }
    injected = append_summary_to_document(doc, summary)
    assert injected is True
    text = "\n".join(p.text for p in doc.paragraphs)
    assert "审计情况小结" in text
    assert "无保留意见" in text

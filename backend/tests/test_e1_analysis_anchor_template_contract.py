"""E1-14 定位锚点表与模板逐格一致（spec wp-deeplink-cell-locate-closure 需求 4.3）。

前端 `e1AnalysisCellAnchors.ts` 登记了模板「货币资金分析表E1-14」三张表的行号 / 行标签 / 表头和三个文字块的行号；
这里读模板文件逐项核对：模板改了行（插行、改名、改表头）而锚点表没跟着改，用例即红。
"""
from __future__ import annotations

import datetime as dt
import re
from pathlib import Path

import openpyxl
import pytest

REPO = Path(__file__).resolve().parents[2]
ANCHORS = REPO / "audit-platform" / "frontend" / "src" / "components" / "workpaper" / "e1" / "e1AnalysisCellAnchors.ts"
SHEET = "货币资金分析表E1-14"


def _template_sheet():
    books = sorted((REPO / "backend" / "wp_templates" / "E").glob("E1-14*.xlsx"))
    assert len(books) == 1, f"E1-14 模板应恰有 1 本，实际 {books}"
    return openpyxl.load_workbook(books[0], data_only=False)[SHEET]


def _parse_tables(src: str) -> list[dict]:
    body = src.split("E1_ANALYSIS_TEMPLATE_TABLES", 1)[1].split("E1_ANALYSIS_TEMPLATE_BLOCKS", 1)[0]
    tables = []
    for chunk in re.split(r"\n\s*\{\s*\n\s*section: ", body)[1:]:
        section = re.match(r"'([a-z]+)'", chunk).group(1)
        header = int(re.search(r"headerRow: (\d+)", chunk).group(1))
        rows = [(int(r), label) for r, label in re.findall(r"\{ row: (\d+), templateLabel: '([^']*)' \}", chunk)]
        cols = dict((c, (html, tpl)) for c, html, tpl in
                    re.findall(r"\b([A-Z]): \{ html: '([^']*)', template: '([^']*)' \}", chunk))
        tables.append({"section": section, "header": header, "rows": rows, "cols": cols})
    return tables


def _parse_blocks(src: str) -> list[dict]:
    body = src.split("E1_ANALYSIS_TEMPLATE_BLOCKS", 1)[1].split("export function", 1)[0]
    return [{"block": b, "rows": [int(x) for x in rows.split(",")], "prefix": prefix}
            for b, rows, prefix in re.findall(
                r"\{ block: '([a-z0-9-]+)', rows: \[([\d, ]+)\][^}]*templatePrefix: '([^']*)' \}", body)]


SRC = ANCHORS.read_text(encoding="utf-8")
TABLES = _parse_tables(SRC)
BLOCKS = _parse_blocks(SRC)


def test_anchor_file_parsed_completely():
    assert [t["section"] for t in TABLES] == ["composition", "other", "ratios"]
    assert [len(t["rows"]) for t in TABLES] == [3, 7, 8]
    assert [len(t["cols"]) for t in TABLES] == [5, 3, 8]
    assert [b["block"] for b in BLOCKS] == ["f3-2", "note", "conclusion"]


@pytest.mark.parametrize("table", TABLES, ids=lambda t: t["section"])
def test_row_labels_match_template_column_b(table):
    ws = _template_sheet()
    mismatched = [(row, label, ws[f"B{row}"].value) for row, label in table["rows"] if ws[f"B{row}"].value != label]
    assert not mismatched, f"锚点表行标签与模板 B 列不一致：{mismatched}"


@pytest.mark.parametrize("table", TABLES, ids=lambda t: t["section"])
def test_column_headers_match_template(table):
    ws = _template_sheet()
    problems = []
    for col, (_, tpl) in table["cols"].items():
        cell = f"{col}{table['header']}"
        if tpl:
            if ws[cell].value != tpl:
                problems.append((cell, tpl, ws[cell].value))
        else:
            merged = next((r for r in ws.merged_cells.ranges if cell in r), None)
            if merged is None or merged.min_row != table["header"] or merged.min_col == ws[cell].column:
                problems.append((cell, "应在表头合并区域内（非左上角）", str(merged)))
    assert not problems, f"锚点表表头与模板不一致：{problems}"


def test_period_columns_are_current_prior_prior2():
    ws = _template_sheet()
    current, prior, prior2 = ws["C13"].value, ws["E13"].value, ws["G13"].value
    assert all(isinstance(v, dt.datetime) for v in (current, prior, prior2)), (current, prior, prior2)
    assert current > prior > prior2


def test_total_rows_are_not_anchored():
    ws = _template_sheet()
    anchored = {row for t in TABLES for row, _ in t["rows"]}
    for row in (18, 29):
        assert ws[f"A{row}"].value == "合计"
        assert row not in anchored, f"第 {row} 行是合计行，网页视图没有对应行，不应登记"


@pytest.mark.parametrize("block", BLOCKS, ids=lambda b: b["block"])
def test_block_rows_start_with_template_text(block):
    ws = _template_sheet()
    first = ws[f"A{block['rows'][0]}"].value
    assert isinstance(first, str) and first.startswith(block["prefix"]), (block, first)

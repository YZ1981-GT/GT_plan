"""汇总导出可回流 + 类型来源优先级 —— Property 6 / 7 / 8。

spec: adjustment-import-export-contract / Task 1.4

Property 6（汇总导出可被导入接受）
  `GET /adjustments/export-summary` 产物的列集合 ⊇ 中央导入必填列集合；
  该产物直接提交导入不因缺列被拒。

Property 7（类型来源优先级）
  「类型」列存在时以列值为准；缺列且 sheet 名可判别时以 sheet 名兜底；
  两者冲突时以列为准并给出提示；两者都无 → 保留「缺少必填列: 类型」错误。

Property 8（导出→导回业务等价）
  汇总导出未修改原样导回后，分录集合业务等价于导出前（编号/类型/金额/科目不失真）。

不依赖真实 DB：export-summary 端点用 fake AsyncSession + AdjustmentService 替身；
导入侧用与 characterization 同款的 fake DB（AccountChart → AccountMapping 两次查询）。
"""
from __future__ import annotations

import asyncio
import io
from decimal import Decimal
from types import SimpleNamespace
from uuid import uuid4

import openpyxl
import pytest

from app.models.audit_platform_models import AccountSource, AdjustmentType
from app.services.import_template_service import (
    TEMPLATE_COLUMNS,
    ImportType,
    infer_adjustment_type_from_sheet_name,
    parse_import_data,
    validate_import_file,
)

_ADJ_COLUMNS = TEMPLATE_COLUMNS[ImportType.adjustments]
_REQUIRED_NAMES = {c[0] for c in _ADJ_COLUMNS if c[1]}

# 汇总导出列头 → 导入侧规范化后的字段名（别名见 `_ADJ_HEADER_ALIASES`）
_SUMMARY_HEADERS = ["编号", "类型", "摘要", "科目编码", "科目名称", "借方金额", "贷方金额", "来源"]


# ═══════════════════════════════════════════════════════════════════════════
# fake DB / service
# ═══════════════════════════════════════════════════════════════════════════

_CHART_ROWS = [
    ("1122", "应收账款", AccountSource.standard),
    ("1231", "坏账准备", AccountSource.standard),
    ("6602", "管理费用", AccountSource.standard),
    ("6701", "信用减值损失", AccountSource.standard),
]
_MAPPING_ROWS: list[tuple[str, str]] = []


class _FakeResult:
    def __init__(self, rows: list):
        self._rows = list(rows)

    def all(self):
        return list(self._rows)

    def scalar_one_or_none(self):
        return self._rows[0] if self._rows else None

    def scalars(self):
        return SimpleNamespace(all=lambda: list(self._rows))


class _FakeDB:
    def __init__(self, results: list[list] | None = None):
        self._queue = [_FakeResult(r) for r in (results or [])]

    async def execute(self, *_a, **_kw):
        if self._queue:
            return self._queue.pop(0)
        return _FakeResult([])

    async def flush(self):
        return None

    async def commit(self):
        return None

    async def rollback(self):
        return None

    def add(self, _obj):
        return None


def _entry(
    adjustment_no: str,
    description: str,
    lines: list[tuple[str, str, str, str]],
    *,
    origin: str = "manual",
) -> dict:
    """构造一个分录组（list_entries 的 item 形状）。"""
    line_items = [
        {
            "standard_account_code": code,
            "account_name": name,
            "debit_amount": debit,
            "credit_amount": credit,
        }
        for code, name, debit, credit in lines
    ]
    return {
        "adjustment_no": adjustment_no,
        "description": description,
        "origin": origin,
        "source_ref": None,
        "line_items": line_items,
        "total_debit": sum(Decimal(li["debit_amount"]) for li in line_items),
        "total_credit": sum(Decimal(li["credit_amount"]) for li in line_items),
    }


_AJE_ENTRIES = [
    _entry("AJE-101", "补提坏账准备", [
        ("6701", "信用减值损失", "120000.50", "0"),
        ("1231", "坏账准备", "0", "120000.50"),
    ]),
    _entry("AJE-102", "费用跨期调整", [
        ("6602", "管理费用", "3000", "0"),
        ("1122", "应收账款", "0", "3000"),
    ], origin="workpaper"),
]
_RJE_ENTRIES = [
    _entry("RJE-201", "应收重分类", [
        ("1122", "应收账款", "88888.88", "0"),
        ("6602", "管理费用", "0", "88888.88"),
    ]),
]


async def _export_summary_bytes(monkeypatch, aje_entries: list[dict], rje_entries: list[dict]) -> bytes:
    """真实调用 `export_adjustment_summary` 端点函数，取回 xlsx 字节。"""
    from app.routers import adjustments as adj_router

    class _Svc:
        def __init__(self, _db):
            pass

        async def list_entries(self, _project_id, _year, adjustment_type=None, **_kw):
            items = aje_entries if adjustment_type == AdjustmentType.aje else rje_entries
            return {"items": items, "total": len(items)}

    monkeypatch.setattr(adj_router, "AdjustmentService", _Svc)

    resp = await adj_router.export_adjustment_summary(
        project_id=uuid4(),
        year=2025,
        format="excel",
        db=_FakeDB(),
        current_user=SimpleNamespace(id=uuid4()),
    )
    chunks: list[bytes] = []
    async for chunk in resp.body_iterator:
        chunks.append(chunk if isinstance(chunk, bytes) else str(chunk).encode())
    return b"".join(chunks)


def _export_summary(monkeypatch, aje=None, rje=None) -> bytes:
    return asyncio.run(
        _export_summary_bytes(
            monkeypatch,
            _AJE_ENTRIES if aje is None else aje,
            _RJE_ENTRIES if rje is None else rje,
        )
    )


# ═══════════════════════════════════════════════════════════════════════════
# Property 6 — 汇总导出可被导入接受
# ═══════════════════════════════════════════════════════════════════════════


def test_p6_summary_export_headers_cover_required_import_columns(monkeypatch):
    """导出列集合（经导入侧表头别名规范化后）⊇ 中央导入必填列集合。"""
    from app.services.import_template_service import normalize_adjustment_header

    content = _export_summary(monkeypatch)
    wb = openpyxl.load_workbook(io.BytesIO(content))
    assert wb.sheetnames == ["AJE审计调整", "RJE重分类"]

    for sheet in wb.sheetnames:
        ws = wb[sheet]
        headers = [ws.cell(row=1, column=c).value for c in range(1, len(_SUMMARY_HEADERS) + 1)]
        assert headers == _SUMMARY_HEADERS, f"{sheet} 汇总列头漂移"
        normalized = {normalize_adjustment_header(str(h)) for h in headers}
        assert _REQUIRED_NAMES <= normalized, (
            f"{sheet} 缺少导入必填列: {sorted(_REQUIRED_NAMES - normalized)}"
        )
    wb.close()


def test_p6_summary_export_passes_import_validation(monkeypatch):
    """导出产物直接送校验：valid=True，不因缺列被拒。"""
    content = _export_summary(monkeypatch)
    result = validate_import_file(ImportType.adjustments, content, "审计调整汇总表_2025.xlsx")

    assert result.valid is True, result.errors
    assert not any("缺少必填列" in e.message for e in result.errors)
    # 「来源」是导出专有列 → 仅作未知列警告，不阻断
    assert any("来源" in w.message for w in result.warnings)


def test_p6_summary_export_parses_without_distortion(monkeypatch):
    """解析导出产物：类型 / 金额 / 科目逐字不失真（两个 sheet 合并读取）。"""
    content = _export_summary(monkeypatch)
    stats: dict = {}
    rows = parse_import_data(ImportType.adjustments, content, stats=stats)

    assert stats["example_skipped"] == 0
    # 无兜底、无冲突时不附带这两个键（与 example_skipped_count 同款「非空才带」）
    assert "type_inferred_from_sheet" not in stats
    assert "type_source_conflicts" not in stats

    parsed = [
        (
            str(r.get("编号")),
            str(r.get("类型")),
            str(r.get("摘要")),
            str(r.get("二级科目编码")),
            str(r.get("二级科目名称")),
            float(r.get("借方金额") or 0),
            float(r.get("贷方金额") or 0),
        )
        for r in rows
    ]
    assert parsed == [
        ("AJE-101", "AJE", "补提坏账准备", "6701", "信用减值损失", 120000.50, 0.0),
        ("AJE-101", "AJE", "补提坏账准备", "1231", "坏账准备", 0.0, 120000.50),
        ("AJE-102", "AJE", "费用跨期调整", "6602", "管理费用", 3000.0, 0.0),
        ("AJE-102", "AJE", "费用跨期调整", "1122", "应收账款", 0.0, 3000.0),
        ("RJE-201", "RJE", "应收重分类", "1122", "应收账款", 88888.88, 0.0),
        ("RJE-201", "RJE", "应收重分类", "6602", "管理费用", 0.0, 88888.88),
    ]


def test_p6_summary_export_group_without_line_items_is_importable(monkeypatch):
    """无明细行的分录组走汇总行路径：类型列同样写入，仍可通过校验。"""
    entry = _entry("AJE-900", "无明细汇总行", [])
    entry["total_debit"] = Decimal("500")
    entry["total_credit"] = Decimal("500")
    content = _export_summary(monkeypatch, aje=[entry], rje=[])

    wb = openpyxl.load_workbook(io.BytesIO(content))
    ws = wb["AJE审计调整"]
    assert [ws.cell(row=2, column=c).value for c in (1, 2, 3, 6, 7)] == [
        "AJE-900", "AJE", "无明细汇总行", 500.0, 500.0,
    ]
    wb.close()

    # 该路径无科目明细 → 只应报「行内必填为空」，不得再报「缺少必填列」
    result = validate_import_file(ImportType.adjustments, content, "t.xlsx")
    assert not any("缺少必填列" in e.message for e in result.errors), result.errors


# ═══════════════════════════════════════════════════════════════════════════
# Property 7 — 类型来源优先级
# ═══════════════════════════════════════════════════════════════════════════


def _build_wb(sheets: list[tuple[str, list[str], list[list]]]) -> bytes:
    """按 (sheet 名, 列头, 数据行) 构造 xlsx。"""
    wb = openpyxl.Workbook()
    wb.remove(wb.active)
    for name, headers, rows in sheets:
        ws = wb.create_sheet(name)
        for col, h in enumerate(headers, 1):
            ws.cell(row=1, column=col, value=h)
        for r, row in enumerate(rows, 2):
            for col, v in enumerate(row, 1):
                ws.cell(row=r, column=col, value=v)
    buf = io.BytesIO()
    wb.save(buf)
    return buf.getvalue()


_WITH_TYPE = ["编号", "类型", "摘要", "科目编码", "科目名称", "借方金额", "贷方金额"]
_WITHOUT_TYPE = ["编号", "摘要", "科目编码", "科目名称", "借方金额", "贷方金额"]


@pytest.mark.parametrize(
    "sheet_name,expected",
    [
        ("AJE审计调整", "AJE"),
        ("RJE重分类", "RJE"),
        ("AJE模板", "AJE"),
        ("RJE模板", "RJE"),
        ("调整模板", None),          # 无关键词 → 不可判别
        ("AJE与RJE合并模板", None),   # 同时命中两类 → 不猜
        ("", None),
    ],
)
def test_p7_sheet_name_inference(sheet_name, expected):
    assert infer_adjustment_type_from_sheet_name(sheet_name) == expected


def test_p7_column_only_is_authoritative():
    """仅有「类型」列（sheet 名不可判别）→ 用列值。"""
    content = _build_wb([
        ("调整模板", _WITH_TYPE, [["RJE-1", "RJE", "仅列", "1122", "应收账款", 100, 0]]),
    ])
    result = validate_import_file(ImportType.adjustments, content, "t.xlsx")
    assert result.valid is True, result.errors

    stats: dict = {}
    rows = parse_import_data(ImportType.adjustments, content, stats=stats)
    assert [r["类型"] for r in rows] == ["RJE"]
    assert "type_inferred_from_sheet" not in stats
    assert "type_source_conflicts" not in stats


def test_p7_sheet_name_fallback_when_column_missing():
    """缺「类型」列但 sheet 名可判别 → 按 sheet 名为该 sheet 全部行注入类型，不再报缺列。"""
    content = _build_wb([
        ("AJE审计调整", _WITHOUT_TYPE, [
            ["AJE-1", "兜底1", "1122", "应收账款", 100, 0],
            ["AJE-1", "兜底1", "1231", "坏账准备", 0, 100],
        ]),
        ("RJE重分类", _WITHOUT_TYPE, [
            ["RJE-1", "兜底2", "6602", "管理费用", 50, 0],
        ]),
    ])

    result = validate_import_file(ImportType.adjustments, content, "t.xlsx")
    assert result.valid is True, result.errors
    assert not any("缺少必填列" in e.message for e in result.errors)
    assert any("按 sheet 名推断" in w.message for w in result.warnings)

    stats: dict = {}
    rows = parse_import_data(ImportType.adjustments, content, stats=stats)
    assert [r["类型"] for r in rows] == ["AJE", "AJE", "RJE"]
    assert stats["type_inferred_from_sheet"] == 3
    assert "type_source_conflicts" not in stats


def test_p7_conflict_column_wins_with_hint():
    """「类型」列与 sheet 名冲突 → 以列为准，并给出冲突提示（不阻断）。"""
    content = _build_wb([
        ("AJE审计调整", _WITH_TYPE, [
            ["X-1", "RJE", "冲突行", "1122", "应收账款", 100, 0],
            ["X-2", "AJE", "一致行", "1231", "坏账准备", 0, 100],
        ]),
    ])

    result = validate_import_file(ImportType.adjustments, content, "t.xlsx")
    assert result.valid is True, result.errors
    conflict_warnings = [w for w in result.warnings if "不一致" in w.message]
    assert len(conflict_warnings) == 1
    assert conflict_warnings[0].row == 2 and conflict_warnings[0].column == "类型"
    assert "以列值为准" in conflict_warnings[0].message

    stats: dict = {}
    rows = parse_import_data(ImportType.adjustments, content, stats=stats)
    assert [r["类型"] for r in rows] == ["RJE", "AJE"], "列值必须胜出"
    conflicts = stats["type_source_conflicts"]
    assert len(conflicts) == 1
    assert conflicts[0] == {
        "sheet": "AJE审计调整", "row": 2,
        "column_value": "RJE", "sheet_inferred": "AJE",
    }


def test_p7_no_column_and_no_sheet_hint_keeps_missing_error():
    """两者都无 → 保留既有「缺少必填列: 类型」错误（明确报错优于静默猜测）。"""
    content = _build_wb([
        ("调整模板", _WITHOUT_TYPE, [["AJE-1", "无从判别", "1122", "应收账款", 100, 0]]),
    ])
    result = validate_import_file(ImportType.adjustments, content, "t.xlsx")

    assert result.valid is False
    assert any("缺少必填列" in e.message and "类型" in e.message for e in result.errors)

    stats: dict = {}
    rows = parse_import_data(ImportType.adjustments, content, stats=stats)
    assert rows and "类型" not in rows[0], "无从判别时不得凭空注入类型"
    assert "type_inferred_from_sheet" not in stats


def test_p7_validate_and_parse_agree_on_multi_sheet_fallback():
    """validate 只看首个 sheet、parse 合并多 sheet → 兜底判定必须一致：

    只要任一候选 sheet 不可判别，校验就不得放行「缺类型」（否则校验过了解析仍缺类型）。
    """
    content = _build_wb([
        ("AJE审计调整", _WITHOUT_TYPE, [["AJE-1", "可判别", "1122", "应收账款", 100, 0]]),
        ("调整模板", _WITHOUT_TYPE, [["AJE-2", "不可判别", "1231", "坏账准备", 0, 100]]),
    ])

    result = validate_import_file(ImportType.adjustments, content, "t.xlsx")
    assert result.valid is False, "第二个 sheet 无从判别类型 → 不得放行"
    assert any("缺少必填列" in e.message for e in result.errors)

    rows = parse_import_data(ImportType.adjustments, content)
    typed = [r for r in rows if r.get("类型")]
    assert len(typed) == 1, "只有可判别的 sheet 被注入类型 —— 正是校验必须拦下的不一致"


# ═══════════════════════════════════════════════════════════════════════════
# Property 8 — 导出→原样导回业务等价
# ═══════════════════════════════════════════════════════════════════════════


class _RecordingSvc:
    """AdjustmentService 替身：记录 create_entry 的入参（导入侧）。"""

    calls: list[tuple] = []

    def __init__(self, _db):
        pass

    async def create_entry(self, project_id, data, user_id):
        _RecordingSvc.calls.append((project_id, data, user_id))
        return SimpleNamespace(id=uuid4(), entry_group_id=uuid4())

    async def delete_entry(self, project_id, entry_group_id):
        return None


class _PinRecordingDB(_FakeDB):
    """在 fake DB 上记录 overwrite 模式的「编号钉住」UPDATE（编号是否原样回流）。

    `AdjustmentCreate` 不含 adjustment_no（编号由 create_entry 自动生成后再被
    `_pin_adjustment_no` 钉为文件编号），故编号等价性只能从该 UPDATE 观察。
    """

    def __init__(self, results: list[list] | None = None):
        super().__init__(results)
        self.pinned: list[str] = []

    async def execute(self, stmt, *a, **kw):
        if "update adjustments" in str(stmt).lower():
            try:
                params = dict(stmt.compile().params)
            except Exception:  # pragma: no cover - 防御
                params = {}
            for key, val in params.items():
                if key == "adjustment_no" or key.startswith("adjustment_no_"):
                    self.pinned.append(val)
            return _FakeResult([])
        return await super().execute(stmt, *a, **kw)


def _business_set(entries: list[dict], adj_type: str) -> set[tuple]:
    """业务等价口径：(类型, 摘要, ((科目码, 借, 贷), ...))。"""
    return {
        (
            adj_type,
            e["description"],
            tuple(
                (
                    li["standard_account_code"],
                    Decimal(str(li["debit_amount"])),
                    Decimal(str(li["credit_amount"])),
                )
                for li in e["line_items"]
            ),
        )
        for e in entries
    }


def _imported_business_set() -> set[tuple]:
    return {
        (
            data.adjustment_type.value.upper(),
            data.description,
            tuple(
                (li.standard_account_code, li.debit_amount, li.credit_amount)
                for li in data.line_items
            ),
        )
        for _pid, data, _uid in _RecordingSvc.calls
    }


def test_p8_export_then_reimport_is_business_equivalent(monkeypatch):
    """汇总导出 → 原样以 overwrite 导回：分录组集合（编号/类型/摘要/科目/借贷）业务等价。"""
    import app.services.adjustment_service as adj_mod
    from app.routers.import_templates import _import_adjustments

    content = _export_summary(monkeypatch)
    rows = parse_import_data(ImportType.adjustments, content)

    _RecordingSvc.calls = []
    monkeypatch.setattr(adj_mod, "AdjustmentService", _RecordingSvc)
    db = _PinRecordingDB([_CHART_ROWS, _MAPPING_ROWS])

    result = asyncio.run(
        _import_adjustments(
            rows, uuid4(), 2025, SimpleNamespace(id=uuid4()), db, mode="overwrite",
        )
    )

    assert result["failed"] == 0, result
    assert result["skipped"] == 0, result
    assert result["imported"] == len(_AJE_ENTRIES) + len(_RJE_ENTRIES)

    # 类型 / 摘要 / 科目 / 借贷金额不失真
    assert _imported_business_set() == (
        _business_set(_AJE_ENTRIES, "AJE") | _business_set(_RJE_ENTRIES, "RJE")
    )
    # 编号原样回流（overwrite 以文件编号为权威键）
    assert sorted(db.pinned) == sorted(
        e["adjustment_no"] for e in (_AJE_ENTRIES + _RJE_ENTRIES)
    )


def test_p8_type_survives_reimport_per_sheet(monkeypatch):
    """RJE 分录导回后仍是 RJE（此前无「类型」列 → 全部退化为 AJE 的失真已消除）。"""
    import app.services.adjustment_service as adj_mod
    from app.routers.import_templates import _import_adjustments

    content = _export_summary(monkeypatch, aje=[], rje=_RJE_ENTRIES)
    rows = parse_import_data(ImportType.adjustments, content)

    _RecordingSvc.calls = []
    monkeypatch.setattr(adj_mod, "AdjustmentService", _RecordingSvc)
    asyncio.run(
        _import_adjustments(
            rows, uuid4(), 2025, SimpleNamespace(id=uuid4()),
            _FakeDB([_CHART_ROWS, _MAPPING_ROWS]),
        )
    )

    assert [d.adjustment_type.value for _p, d, _u in _RecordingSvc.calls] == ["rje"]


def test_p8_sheet_name_only_export_would_lose_rje_without_type_column():
    """反例锚点：同一批数据若无「类型」列且 sheet 名不可判别，RJE 会退化为 AJE。

    说明本 task 补「类型」列 + sheet 名兜底的必要性（Property 7/8 的动机）。
    """
    content = _build_wb([
        ("调整模板", _WITHOUT_TYPE, [
            ["RJE-201", "应收重分类", "1122", "应收账款", 88888.88, 0],
        ]),
    ])
    rows = parse_import_data(ImportType.adjustments, content)
    assert "类型" not in rows[0]
    # `_import_adjustments` 对缺类型的行默认按 AJE 处理 → 失真（故必须有类型来源）
    assert str(rows[0].get("类型", "") or "AJE").upper() == "AJE"

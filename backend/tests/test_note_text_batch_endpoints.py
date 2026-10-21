"""附注正文批量导入端点（R-01 / R-18）：TestClient 真发请求，鉴权与依赖链真实执行。

只覆盖 ``get_db`` / ``get_current_user``（``require_project_access`` 每次返回新函数，不能直接覆盖）；
协调器 Redis 打桩为不可用、EventBus 只记录不分发。SQLite 内存库 + 真 ORM。
"""

from __future__ import annotations

import hashlib
import io
import uuid
from datetime import datetime, timezone
from collections.abc import AsyncGenerator
from unittest.mock import AsyncMock, patch

import pytest
import pytest_asyncio
import sqlalchemy as sa
from fastapi import FastAPI
from httpx import ASGITransport, AsyncClient
from sqlalchemy.ext.asyncio import AsyncSession, async_sessionmaker, create_async_engine

import tests.conftest  # noqa: F401  注册全部模型 + SQLite JSONB 兼容
from app.core.database import get_db
from app.deps import get_current_user
from app.models.base import Base, PermissionLevel, ProjectStatus, ProjectUserRole
from app.models.core import Project, ProjectUser, User, UserRole
from app.models.knowledge_models import KnowledgeAccessLevel, KnowledgeDocument, KnowledgeFolder
from app.models.note_content_template_models import (
    NoteContentTemplate,
    NoteContentTemplateVersion,
    NoteTextBatchApplication,
)
from app.models.report_models import DisclosureNote
from app.models.staff_models import ProjectAssignment, StaffMember
from app.routers import note_text_batch
from app.services.disclosure_mutation_coordinator import DisclosureMutationCoordinator

YEAR = 2025
# 来源端点要解析写入人姓名（人员档案 → 登录名），与生产一样需要 staff_members
_TABLES = [User.__table__, Project.__table__, ProjectUser.__table__, DisclosureNote.__table__, StaffMember.__table__,
           ProjectAssignment.__table__,
           NoteContentTemplate.__table__, NoteContentTemplateVersion.__table__, NoteTextBatchApplication.__table__,
           KnowledgeFolder.__table__, KnowledgeDocument.__table__]
BASE = "/api/projects/{pid}"


class _Actor:
    def __init__(self, uid: uuid.UUID) -> None:
        self.id, self.role, self.is_active = uid, UserRole.auditor, True
        self.username = f"u-{uid.hex[-4:]}"
        self.email = f"{self.username}@example.com"


def _docx(*blocks: tuple) -> bytes:
    """``(标题, 正文[, 标题级别=2])``；正文为空不写段落。"""
    import docx

    doc = docx.Document()
    for heading, body, *level in blocks:
        doc.add_heading(heading, level=level[0] if level else 2)
        if body:
            doc.add_paragraph(body)
    buf = io.BytesIO()
    doc.save(buf)
    return buf.getvalue()


#: 按模板树形写的上传文档（第 23 组按章节树逐级定位）：一级章「公司基本情况」自带正文、「重要会计政策」下有「会计期间」
TREE_DOC = (("公司基本情况", "本公司成立于2001年。", 1), ("重要会计政策、会计估计", "", 1), ("会计期间", "自公历1月1日起。", 2))


@pytest_asyncio.fixture
async def env() -> AsyncGenerator[dict, None]:
    engine = create_async_engine("sqlite+aiosqlite:///:memory:")
    async with engine.begin() as conn:
        await conn.run_sync(lambda c: Base.metadata.create_all(c, tables=_TABLES))
    factory = async_sessionmaker(engine, class_=AsyncSession, expire_on_commit=False)
    ids = {k: uuid.uuid4() for k in ("me", "peer", "viewer", "outsider", "other", "p", "src", "basic", "period",
                                      "cash", "foreign", "pub", "priv", "doc", "secret", "tpl", "ver", "tpl_x", "ver_x")}
    async with factory() as s:
        s.add_all([
            Project(id=ids["p"], name="甲_2025", client_name="甲公司", template_type="soe",
                    report_scope="standalone", status=ProjectStatus.execution),
            Project(id=ids["src"], name="乙_2025", client_name="乙公司", template_type="soe",
                    report_scope="standalone", status=ProjectStatus.execution),
        ])
        await s.flush()
        s.add_all([
            ProjectUser(project_id=ids["p"], user_id=ids["me"], role=ProjectUserRole.auditor,
                        permission_level=PermissionLevel.edit),
            ProjectUser(project_id=ids["p"], user_id=ids["peer"], role=ProjectUserRole.auditor,
                        permission_level=PermissionLevel.edit),
            ProjectUser(project_id=ids["p"], user_id=ids["viewer"], role=ProjectUserRole.auditor),
            # 章节树（section_id / parent_section_id，与生成链 / Task 10 回填同形）：一级章是容器，正文在其下节点
            *[DisclosureNote(project_id=ids["p"], year=YEAR, note_section=code, section_title=title,
                             section_id=sid, level=1, sort_index=0, sort_order=order,
                             table_data=None)  # 显式 None ⇒ JSON null（与生成链一致，见 test_note_text_batch_targets）
              for code, title, sid, order in (("一", "公司基本情况", "ch-01", 0), ("四", "重要会计政策、会计估计", "ch-04", 10),
                                               ("八", "财务报表主要项目注释", "ch-08", 20))],
            DisclosureNote(id=ids["basic"], project_id=ids["p"], year=YEAR, note_section="一、1",
                           section_title="公司基本情况", text_content="<p>旧基本情况</p>", sort_index=1,
                           section_id="s-01-01", parent_section_id="ch-01", level=2, sort_order=1),
            DisclosureNote(id=ids["period"], project_id=ids["p"], year=YEAR, note_section="四、会计期间",
                           section_title="会计期间", sort_index=2,
                           section_id="s-04-01", parent_section_id="ch-04", level=2, sort_order=11),
            DisclosureNote(id=ids["cash"], project_id=ids["p"], year=YEAR, note_section="八、1",
                           section_title="货币资金", text_content="<p>货币资金说明</p>", sort_index=3,
                           section_id="s-08-01", parent_section_id="ch-08", level=2, sort_order=21),
            DisclosureNote(id=ids["foreign"], project_id=ids["src"], year=YEAR, note_section="四、会计期间",
                           section_title="会计期间", sort_index=2),
            KnowledgeFolder(id=ids["pub"], name="公开", access_level=KnowledgeAccessLevel.public, created_by=ids["other"]),
            KnowledgeFolder(id=ids["priv"], name="他人私有", access_level=KnowledgeAccessLevel.private,
                            created_by=ids["other"]),
        ])
        await s.flush()
        s.add_all([
            KnowledgeDocument(id=ids["doc"], folder_id=ids["pub"], name="政策汇编.md",
                              content_text="# 公司基本情况\n本公司由知识库预填。\n# 重要会计政策、会计估计\n"
                                           "## 会计期间\n自公历1月1日起。"),
            KnowledgeDocument(id=ids["secret"], folder_id=ids["priv"], name="私密底稿.md",
                              content_text="# 公司基本情况\n机密正文不得外泄"),
            *_template_rows(ids),
        ])
        await s.commit()

    current = {"actor": _Actor(ids["me"])}
    app = FastAPI()
    app.include_router(note_text_batch.router)

    async def _override_db():
        async with factory() as session:
            yield session

    app.dependency_overrides[get_db] = _override_db
    app.dependency_overrides[get_current_user] = lambda: current["actor"]
    published: list = []
    with patch.object(DisclosureMutationCoordinator, "_redis", new=AsyncMock(return_value=None)), \
            patch("app.services.event_bus.event_bus.publish", new=AsyncMock(side_effect=published.append)):
        async with AsyncClient(transport=ASGITransport(app=app), base_url="http://test") as client:
            yield {"client": client, "ids": ids, "factory": factory, "published": published,
                   "act_as": lambda key: current.update(actor=_Actor(ids[key]))}
    await engine.dispose()


DOCX_MIME = "application/vnd.openxmlformats-officedocument.wordprocessingml.document"
TPL_HASH = "a" * 64


def _template_rows(ids: dict) -> list:
    def version(vid: uuid.UUID, tid: uuid.UUID, text: str) -> NoteContentTemplateVersion:
        return NoteContentTemplateVersion(
            id=vid, template_id=tid, version_no=1, source_project_id=ids["src"], source_year=YEAR,
            source_project_name="乙_2025", template_type="soe", report_scope="standalone",
            sections_data=[{"source_note_id": str(uuid.uuid4()), "section_id": None, "note_section": "一、1",
                            "section_title": "公司基本情况", "section_kind": "narrative", "content_type": "text",
                            "text_content": text, "content_hash": "0" * 64, "status": "confirmed",
                            "is_stale": False, "quality_flags": []}],
            quality_summary={}, content_hash=TPL_HASH, byte_size=len(text))

    return [NoteContentTemplate(id=ids["tpl"], name="乙公司全套", owner_scope="personal", owner_user_id=ids["me"]),
            NoteContentTemplate(id=ids["tpl_x"], name="他人模板", owner_scope="personal", owner_user_id=ids["other"]),
            version(ids["ver"], ids["tpl"], "<p>乙公司成立于2001年。</p>"),
            version(ids["ver_x"], ids["tpl_x"], "<p>他人模板正文</p>")]


def _url(env: dict, tail: str) -> str:
    return BASE.format(pid=env["ids"]["p"]) + tail


async def _add_version(env: dict, sections: list[tuple[str, str, str, str]], *, template_type: str = "soe") -> dict:
    """给本人模板追加一个版本，返回可直接用于 preview 的模板来源。"""
    vid, content_hash = uuid.uuid4(), uuid.uuid4().hex * 2
    async with env["factory"]() as db:
        db.add(NoteContentTemplateVersion(
            id=vid, template_id=env["ids"]["tpl"], version_no=9, source_project_id=env["ids"]["src"],
            source_year=YEAR, source_project_name="乙_2025", template_type=template_type, report_scope="standalone",
            sections_data=[{"source_note_id": str(uuid.uuid4()), "note_section": code, "section_title": title,
                            "section_kind": kind, "text_content": html, "status": "confirmed", "is_stale": False}
                           for code, title, kind, html in sections],
            quality_summary={}, content_hash=content_hash, byte_size=1))
        await db.commit()
    return {"source_kind": "content_template", "template_version_id": str(vid), "template_content_hash": content_hash}


async def _parse(env: dict, data: bytes | None = None, *, name: str = "附注.docx", ctype: str = DOCX_MIME):
    data = data if data is not None else _docx(*TREE_DOC)
    return await env["client"].post(_url(env, "/disclosure-note-import/parse-upload"),
                                    files={"file": (name, data, ctype)})


async def _upload_source(env: dict) -> dict:
    resp = await _parse(env)
    assert resp.status_code == 200, resp.text
    return {"source_kind": "upload", "parsed_source": resp.json()}


async def _preview(env: dict, source: dict, **extra):
    return await env["client"].post(_url(env, "/disclosure-note-text-batch/preview"),
                                    json={"year": YEAR, "source": source, **extra})


def _item(preview: dict, title: str) -> dict:
    return next(i for i in preview["items"] if i["source_title"] == title)


def _assert_error(resp, status: int, code: str) -> dict:
    """先断言状态码（红时打印响应体，能看出是「被放行」还是「被别的校验拦下」），再断言错误码。"""
    assert resp.status_code == status, resp.text
    detail = resp.json()["detail"]
    assert detail["error_code"] == code, detail
    return detail


def _kb_body(doc_ids: list[str], section_doc: str) -> dict:
    return {"year": YEAR, "source": {"source_kind": "knowledge", "knowledge_doc_ids": doc_ids, "parsed_source": {
        "source_kind": "knowledge", "sections": [{"source_id": f"{section_doc}:s1", "title": "公司基本情况",
                                                  "html": "<p>x</p>", "metadata": {"doc_id": section_doc}}]}}}


def _selection(item: dict, mode: str = "replace") -> dict:
    return {"source_id": item["source_id"], "target_note_id": item["target_note_id"], "mode": mode,
            "before_hash": item["before_hash"], "expected_updated_at": item["expected_updated_at"],
            "proposed_hash": item["append_hash"] if mode == "append" else item["proposed_hash"],
            "proposed_html": "<p>浏览器提交的正文一律忽略</p>"}


async def _apply(env: dict, source: dict, preview: dict, selections: list[dict], *,
                 mutation_id: str | None = "mut-000001", **extra):
    body = {"year": YEAR, "source": source, "preview_digest": preview["preview_digest"],
            "selections": selections, **extra}
    headers = {"X-Mutation-Id": mutation_id} if mutation_id else {}
    return await env["client"].post(_url(env, "/disclosure-note-text-batch/apply"), json=body, headers=headers)


async def _note(env: dict, key: str) -> DisclosureNote:
    async with env["factory"]() as db:
        return await db.get(DisclosureNote, env["ids"][key])


async def _applications(env: dict) -> list:
    async with env["factory"]() as db:
        return (await db.execute(NoteTextBatchApplication.__table__.select())).all()


async def _update_project(env: dict, **values) -> None:
    async with env["factory"]() as db:
        project = await db.get(Project, env["ids"]["p"])
        for key, value in values.items():
            setattr(project, key, value)
        await db.commit()


class TestSections:
    @pytest.mark.asyncio
    async def test_readonly_member_lists_narrative_without_content(self, env):
        env["act_as"]("viewer")
        resp = await env["client"].get(_url(env, "/disclosure-note-text-sections"), params={"year": YEAR})
        assert resp.status_code == 200, resp.text
        items = resp.json()["items"]
        assert {i["note_section"] for i in items} == {"一、1", "四、会计期间"}
        assert all("text_content" not in i for i in items)
        full = await env["client"].get(_url(env, "/disclosure-note-text-sections"), params={"year": YEAR, "scope": "all"})
        assert "八、1" in {i["note_section"] for i in full.json()["items"]}

    @pytest.mark.asyncio
    async def test_excludes_containers_deleted_and_out_of_scope_sections(self, env):
        """Task 21：narrative 不含容器；all 标出容器；软删除与口径不适用（单体项目里的合并专属章）两者都排除。"""
        from app.services.note_section_catalog import load_section_scope_map, section_applies_to_scope

        consol_only = sorted(c for c, s in load_section_scope_map("soe").items()
                             if not section_applies_to_scope({"scope": s}, "standalone"))[0]
        pid = env["ids"]["p"]
        async with env["factory"]() as db:
            db.add_all([
                DisclosureNote(project_id=pid, year=YEAR, note_section="四、2", section_title="重要会计政策",
                               section_id="sec-policy", sort_index=10),
                DisclosureNote(project_id=pid, year=YEAR, note_section="四、2、1", section_title="收入确认",
                               parent_section_id="sec-policy", text_content="<p>收入</p>", sort_index=11),
                DisclosureNote(project_id=pid, year=YEAR, note_section="四、3", section_title="已删章节",
                               text_content="<p>x</p>", is_deleted=True, sort_index=12),
                DisclosureNote(project_id=pid, year=YEAR, note_section=consol_only, section_title="合并专属",
                               text_content="<p>x</p>", sort_index=13),
            ])
            await db.commit()

        async def _sections(scope: str) -> dict[str, str]:
            resp = await env["client"].get(_url(env, "/disclosure-note-text-sections"), params={"year": YEAR, "scope": scope})
            return {i["note_section"]: i["section_kind"] for i in resp.json()["items"]}

        narrative, every = await _sections("narrative"), await _sections("all")
        assert "四、2、1" in narrative and "四、2" not in narrative
        assert every["四、2"] == "container"
        assert "四、3" not in every, "软删除章节被列出"
        assert consol_only not in every, "单体项目列出了合并专属章节（口径不适用）"

    @pytest.mark.asyncio
    async def test_outsider_forbidden(self, env):
        env["act_as"]("outsider")
        resp = await env["client"].get(_url(env, "/disclosure-note-text-sections"), params={"year": YEAR})
        assert resp.status_code == 403


class TestParseUpload:
    @pytest.mark.asyncio
    async def test_docx_parsed_and_metadata_signed_by_server(self, env):
        import hashlib

        data = _docx(("公司基本情况", "本公司成立于2001年。"))
        resp = await _parse(env, data)
        assert resp.status_code == 200, resp.text
        body = resp.json()
        assert "公司基本情况" in [s["title"] for s in body["sections"]]
        meta = body["source_metadata"]
        assert (meta["file_hash"], meta["byte_size"], meta["file_format"]) == (hashlib.sha256(data).hexdigest(),
                                                                               len(data), "docx")
        assert len(meta["signature"]) == 64

    @pytest.mark.parametrize("name,data,ctype,reason", [
        ("附注.txt", b"plain text", "text/plain", "仅支持"),                          # 扩展名
        ("附注.docx", b"%PDF-1.4 pretend docx", DOCX_MIME, "签名"),                 # 签名与扩展名不符
        ("附注.pdf", b"PK\x03\x04 pretend pdf", "application/pdf", "签名"),         # 签名与扩展名不符
        ("附注.docx", None, "text/html", "类型与扩展名不符"),                         # MIME 与扩展名不符
    ])
    @pytest.mark.asyncio
    async def test_extension_mime_signature_all_checked(self, env, name, data, ctype, reason):
        resp = await _parse(env, data if data is not None else _docx(("公司基本情况", "正文")), name=name, ctype=ctype)
        assert resp.status_code == 422, resp.text
        detail = resp.json()["detail"]
        # 解析器对坏文件同样报 UNSUPPORTED_FILE ⇒ 必须按原因区分，才能证明是前置校验拦下的
        assert detail["error_code"] == "UNSUPPORTED_FILE" and reason in detail["message"]

    @pytest.mark.asyncio
    async def test_text_pdf_parsed(self, env):
        import pymupdf

        doc = pymupdf.open()
        page = doc.new_page()
        page.insert_text((72, 72), "一、公司基本情况", fontsize=20, fontname="china-s")
        page.insert_text((72, 110), "本公司成立于2001年。", fontsize=10, fontname="china-s")
        data = doc.tobytes()
        doc.close()
        resp = await _parse(env, data, name="附注.pdf", ctype="application/pdf")
        assert resp.status_code == 200, resp.text
        assert resp.json()["sections"] and resp.json()["source_metadata"]["file_format"] == "pdf"

    @pytest.mark.parametrize("kind,code", [("scanned_pdf", "PDF_NEEDS_OCR"), ("empty_docx", "NO_PARSEABLE_TEXT")])
    @pytest.mark.asyncio
    async def test_no_text_is_explained_not_empty_success(self, env, kind, code):
        """AC-1.5：扫描件 / 空文档给中文原因，不返回「解析成功但内容为空」。"""
        import docx
        import pymupdf

        if kind == "scanned_pdf":
            doc = pymupdf.open()
            doc.new_page()
            data, name, ctype = doc.tobytes(), "扫描件.pdf", "application/pdf"
            doc.close()
        else:
            buf = io.BytesIO()
            docx.Document().save(buf)
            data, name, ctype = buf.getvalue(), "空.docx", DOCX_MIME
        _assert_error(await _parse(env, data, name=name, ctype=ctype), 422, code)

    @pytest.mark.asyncio
    async def test_oversize_rejected_413(self, env):
        with patch.object(note_text_batch, "MAX_UPLOAD_BYTES", 1000), \
                patch("app.services.note_text_batch_flow.MAX_UPLOAD_BYTES", 1000):
            resp = await _parse(env)
        assert resp.status_code == 413
        assert resp.json()["detail"]["error_code"] == "FILE_TOO_LARGE"

    @pytest.mark.asyncio
    async def test_readonly_member_cannot_parse(self, env):
        env["act_as"]("viewer")
        assert (await _parse(env)).status_code == 403


class TestParseKnowledge:
    @pytest.mark.asyncio
    async def test_visible_doc_parsed_invisible_doc_hidden(self, env):
        ids = env["ids"]
        resp = await env["client"].post(_url(env, "/disclosure-note-import/parse-knowledge"),
                                        json={"year": YEAR, "doc_ids": [str(ids["doc"]), str(ids["secret"])]})
        assert resp.status_code == 200, resp.text
        body = resp.json()
        assert body["sections"] and all(s["source_id"].startswith(f"{ids['doc']}:") for s in body["sections"])
        # 章节树上级按文档命名空间（AC-1.6）：「会计期间」挂在「重要会计政策、会计估计」下
        by_title = {s["title"]: s for s in body["sections"]}
        assert by_title["会计期间"]["parent_source_id"] == by_title["重要会计政策、会计估计"]["source_id"]
        assert by_title["公司基本情况"]["parent_source_id"] is None
        assert any("1 篇文档不可用" in w for w in body["warnings"])
        assert "机密正文" not in resp.text and "私密底稿" not in resp.text

    @pytest.mark.asyncio
    async def test_only_invisible_doc_is_422_without_leak(self, env):
        resp = await env["client"].post(_url(env, "/disclosure-note-import/parse-knowledge"),
                                        json={"year": YEAR, "doc_ids": [str(env["ids"]["secret"])]})
        assert resp.status_code == 422
        assert resp.json()["detail"]["error_code"] == "NO_PARSEABLE_TEXT"
        assert "机密正文" not in resp.text and "私密底稿" not in resp.text


class TestPreview:
    @pytest.mark.asyncio
    async def test_upload_preview_sanitized_on_server(self, env):
        source = await _upload_source(env)
        section = next(s for s in source["parsed_source"]["sections"] if s["title"] == "公司基本情况")
        section["html"] = '<script>alert(1)</script><p onclick="steal()">恶意正文</p><img src=x onerror=y>'
        resp = await _preview(env, source)
        assert resp.status_code == 200, resp.text
        proposed = _item(resp.json(), "公司基本情况")["proposed_html"]
        assert proposed == "<p>恶意正文</p>"

    @pytest.mark.asyncio
    async def test_tampered_upload_metadata_rejected(self, env):
        source = await _upload_source(env)
        source["parsed_source"]["source_metadata"]["file_hash"] = "0" * 64
        resp = await _preview(env, source)
        assert resp.status_code == 422
        assert resp.json()["detail"]["error_code"] == "SOURCE_METADATA_INVALID"

    @pytest.mark.asyncio
    async def test_upload_metadata_bound_to_parsing_user(self, env):
        """签名绑定解析人：同项目另一位编辑者拿别人的解析结果去预览 ⇒ 拒绝。"""
        source = await _upload_source(env)
        env["act_as"]("peer")
        _assert_error(await _preview(env, source), 422, "SOURCE_METADATA_INVALID")

    @pytest.mark.parametrize("endpoint", ["parse-knowledge", "preview"])
    @pytest.mark.asyncio
    async def test_readonly_member_cannot_use_write_flow(self, env, endpoint):
        # 公开知识库文档对只读成员可见：去掉 edit 门禁后请求会真成功，红的形态即「被放行」
        doc = str(env["ids"]["doc"])
        path, body = (("/disclosure-note-import/parse-knowledge", {"year": YEAR, "doc_ids": [doc]})
                      if endpoint == "parse-knowledge" else ("/disclosure-note-text-batch/preview", _kb_body([doc], doc)))
        env["act_as"]("viewer")
        resp = await env["client"].post(_url(env, path), json=body)
        assert resp.status_code == 403, resp.text

    @pytest.mark.asyncio
    async def test_foreign_project_mapping_rejected(self, env):
        source = await _upload_source(env)
        sid = source["parsed_source"]["sections"][0]["source_id"]
        resp = await _preview(env, source, manual_mappings=[{"source_id": sid,
                                                             "target_note_id": str(env["ids"]["foreign"])}])
        assert resp.status_code == 422
        assert resp.json()["detail"]["error_code"] == "MAPPING_INVALID"

    @pytest.mark.asyncio
    async def test_oversized_section_rejected_413(self, env):
        source = await _upload_source(env)
        source["parsed_source"]["sections"][0]["html"] = "<p>" + "甲" * 200_000 + "</p>"  # 600KB > 单章 512KB
        resp = await _preview(env, source)
        assert resp.status_code == 413
        assert resp.json()["detail"]["error_code"] == "PAYLOAD_TOO_LARGE"

    @pytest.mark.asyncio
    async def test_template_preview_pins_version_and_replaces_source_entity(self, env):
        source = {"source_kind": "content_template", "template_version_id": str(env["ids"]["ver"]),
                  "template_content_hash": TPL_HASH}
        resp = await _preview(env, source)
        assert resp.status_code == 200, resp.text
        item = next(i for i in resp.json()["items"] if i["target_note_section"] == "一、1")
        assert item["confidence"] == "confirmed"
        assert item["proposed_html"] == "<p>甲公司成立于2001年。</p>"
        assert [r["old_value"] for r in item["replacements"]] == ["乙公司"]

    @pytest.mark.asyncio
    async def test_template_groups_account_explanation_as_review(self, env):
        """AC-9.1 / 9.2：确定映射的空叙述章节推荐并默认勾选；报表科目解释章节只能进「需复核」。"""
        async with env["factory"]() as db:
            (await db.get(DisclosureNote, env["ids"]["cash"])).text_content = None
            await db.commit()
        source = await _add_version(env, [("四、会计期间", "会计期间", "narrative", "<p>本公司会计年度采用公历年度。</p>"),
                                           ("八、1", "货币资金", "account_explanation", "<p>货币资金包括库存现金。</p>")])
        resp = await _preview(env, source)
        assert resp.status_code == 200, resp.text
        by_code = {i["target_note_section"]: i for i in resp.json()["items"]}
        assert (by_code["四、会计期间"]["category"], by_code["四、会计期间"]["default_selected"]) == ("recommended", True)
        cash = by_code["八、1"]
        assert (cash["category"], cash["default_selected"]) == ("review", False)
        assert "account_explanation" in cash["quality_flags"]

    @pytest.mark.asyncio
    async def test_template_from_other_standard_flagged(self, env):
        """AC-8.5：模板标准与目标不一致 ⇒ 醒目标记并进「需复核」。"""
        source = await _add_version(env, [("四、会计期间", "会计期间", "narrative", "<p>本公司会计年度采用公历年度。</p>")],
                                    template_type="listed")
        items = [i for i in (await _preview(env, source)).json()["items"] if i["target_note_id"]]
        assert items and all("cross_standard" in i["quality_flags"] and i["category"] == "review" for i in items)

    @pytest.mark.asyncio
    async def test_upload_titles_are_only_candidates(self, env):
        """AC-3.3：外部文档按标题匹配只是候选，不进推荐、不默认勾选。"""
        items = (await _preview(env, await _upload_source(env))).json()["items"]
        matched = [i for i in items if i["target_note_id"]]
        assert matched and all(i["confidence"] == "candidate" and not i["default_selected"] for i in matched)

    @pytest.mark.parametrize("version_key,content_hash,status,code", [
        ("ver_x", TPL_HASH, 404, "TEMPLATE_NOT_VISIBLE"),        # 他人个人模板
        ("ver", "b" * 64, 409, "TEMPLATE_HASH_MISMATCH"),        # 固定 hash 对不上
    ])
    @pytest.mark.asyncio
    async def test_template_visibility_and_pinned_hash(self, env, version_key, content_hash, status, code):
        source = {"source_kind": "content_template", "template_version_id": str(env["ids"][version_key]),
                  "template_content_hash": content_hash}
        resp = await _preview(env, source)
        assert resp.status_code == status, resp.text
        assert resp.json()["detail"]["error_code"] == code
        assert "他人模板正文" not in resp.text

    @pytest.mark.parametrize("claimed,section_doc,status,code", [
        # 声称的文档里有不可见的（章节只来自可见文档）⇒ 文档集合复核拦下
        (("doc", "secret"), "doc", 404, "KNOWLEDGE_DOC_NOT_VISIBLE"),
        # 声称的文档都可见，但章节挂在别的文档上 ⇒ 章节归属复核拦下
        (("doc",), "secret", 422, "SOURCE_SECTION_INVALID"),
    ])
    @pytest.mark.asyncio
    async def test_knowledge_doc_rechecked_on_preview(self, env, claimed, section_doc, status, code):
        body = _kb_body([str(env["ids"][k]) for k in claimed], str(env["ids"][section_doc]))
        resp = await env["client"].post(_url(env, "/disclosure-note-text-batch/preview"), json=body)
        _assert_error(resp, status, code)


@pytest.mark.asyncio
async def test_client_section_metadata_is_discarded():
    """章类型 / 编号元数据只信服务端：客户端回传的一律丢弃（只留 doc_id）。"""
    from app.models.note_content_template_schemas import ParsedSource
    from app.services.note_text_batch_flow import _client_sections

    parsed = ParsedSource.model_validate({"source_kind": "upload", "sections": [
        {"source_id": "s1", "title": "货币资金", "html": "<p>x</p>",
         "metadata": {"section_kind": "narrative", "note_section": "八、1", "doc_id": "d-1"}}]})
    assert _client_sections(parsed, "upload")[0].metadata == {"doc_id": "d-1"}


async def _knowledge_source(env: dict) -> dict:
    doc = str(env["ids"]["doc"])
    resp = await env["client"].post(_url(env, "/disclosure-note-import/parse-knowledge"),
                                    json={"year": YEAR, "doc_ids": [doc]})
    assert resp.status_code == 200, resp.text
    return {"source_kind": "knowledge", "knowledge_doc_ids": [doc], "parsed_source": resp.json()}


async def _previewed(env: dict, kind: str = "upload") -> tuple[dict, dict, dict]:
    source = await (_upload_source(env) if kind == "upload" else _knowledge_source(env))
    resp = await _preview(env, source)
    assert resp.status_code == 200, resp.text
    preview = resp.json()
    return source, preview, _item(preview, "公司基本情况")


class TestApply:
    @pytest.mark.asyncio
    async def test_writes_server_preview_records_audit_and_publishes_event(self, env):
        from app.models.audit_platform_schemas import EventType

        source, preview, item = await _previewed(env)
        resp = await _apply(env, source, preview, [_selection(item)])
        assert resp.status_code == 200, resp.text
        assert resp.json()["applied_count"] == 1
        note = await _note(env, "basic")
        assert note.text_content == item["proposed_html"] == "<p>本公司成立于2001年。</p>"
        assert note.is_local_override is True
        [event] = env["published"]
        assert (event.event_type, event.extra["section_code"]) == (EventType.NOTE_SECTION_SAVED, "一、1")
        [row] = await _applications(env)
        assert (row.source_kind, row.mutation_id) == ("upload", "mut-000001")
        assert row.source_metadata["file_hash"] == source["parsed_source"]["source_metadata"]["file_hash"]
        assert "signature" not in row.source_metadata

    @pytest.mark.asyncio
    async def test_same_mutation_replays_and_different_payload_conflicts(self, env):
        source, preview, item = await _previewed(env)
        first = await _apply(env, source, preview, [_selection(item)])
        again = await _apply(env, source, preview, [_selection(item)])
        assert again.status_code == 200 and again.json()["application_id"] == first.json()["application_id"]
        assert len(await _applications(env)) == 1 and len(env["published"]) == 1
        other = await _apply(env, source, preview, [_selection(item, "append")])
        _assert_error(other, 409, "MUTATION_ID_CONFLICT")

    @pytest.mark.asyncio
    async def test_second_mutation_on_same_targets_loses_with_409(self, env):
        """两个不同 mutation 用同一份预览：先到者写入，后到者按逐章 before_hash 判 stale、零写入。"""
        source, preview, item = await _previewed(env)
        assert (await _apply(env, source, preview, [_selection(item)], mutation_id="mut-first-1")).status_code == 200
        second = await _apply(env, source, preview, [_selection(item)], mutation_id="mut-second-1")
        _assert_error(second, 409, "NOTE_BATCH_STALE_PREVIEW")
        assert len(await _applications(env)) == 1 and len(env["published"]) == 1

    @pytest.mark.asyncio
    async def test_new_template_version_after_preview_does_not_change_result(self, env):
        """预览后模板发布新版本：apply 仍按预览固定的 version_id + content_hash 写入（AC-8.3 / 8.4）。"""
        ids = env["ids"]
        source = {"source_kind": "content_template", "template_version_id": str(ids["ver"]),
                  "template_content_hash": TPL_HASH}
        preview = (await _preview(env, source)).json()
        item = next(i for i in preview["items"] if i["target_note_section"] == "一、1")
        async with env["factory"]() as db:
            newer = NoteContentTemplateVersion(
                id=uuid.uuid4(), template_id=ids["tpl"], version_no=2, source_project_id=ids["src"],
                source_year=YEAR, source_project_name="乙_2025", template_type="soe", report_scope="standalone",
                sections_data=[{"note_section": "一、1", "section_title": "公司基本情况", "section_kind": "narrative",
                                "text_content": "<p>新版本正文</p>"}],
                quality_summary={}, content_hash="c" * 64, byte_size=10)
            db.add(newer)
            await db.flush()
            (await db.get(NoteContentTemplate, ids["tpl"])).current_version_id = newer.id
            await db.commit()
        resp = await _apply(env, source, preview, [_selection(item)])
        assert resp.status_code == 200, resp.text
        assert (await _note(env, "basic")).text_content == "<p>甲公司成立于2001年。</p>"
        [row] = await _applications(env)
        assert (uuid.UUID(str(row.template_version_id)), row.source_metadata["content_hash"]) == (ids["ver"], TPL_HASH)
        lineage = (await _note(env, "basic")).template_lineage["content_templates"][-1]
        assert lineage["template_version_id"] == str(ids["ver"])

    @pytest.mark.asyncio
    async def test_new_version_reaches_body_only_after_user_repreviews_and_applies(self, env):
        """AC-8.4 / Task 51：已引用 v1 的项目在模板发布 v2 后正文不变（只预览也不变）；
        用户用 v2 重新预览并确认写入后才变，两条应用记录各自钉住所用版本。"""
        ids = env["ids"]
        v1_source = {"source_kind": "content_template", "template_version_id": str(ids["ver"]),
                     "template_content_hash": TPL_HASH}
        v1_preview = (await _preview(env, v1_source)).json()
        v1_item = next(i for i in v1_preview["items"] if i["target_note_section"] == "一、1")
        resp = await _apply(env, v1_source, v1_preview, [_selection(v1_item)], mutation_id="mut-v1-000001")
        assert resp.status_code == 200, resp.text
        v1_body = (await _note(env, "basic")).text_content
        async with env["factory"]() as db:
            newer = NoteContentTemplateVersion(
                id=uuid.uuid4(), template_id=ids["tpl"], version_no=2, source_project_id=ids["src"],
                source_year=YEAR, source_project_name="乙_2025", template_type="soe", report_scope="standalone",
                sections_data=[{"note_section": "一、1", "section_title": "公司基本情况", "section_kind": "narrative",
                                "text_content": "<p>新版本正文</p>"}],
                quality_summary={}, content_hash="c" * 64, byte_size=10)
            db.add(newer)
            await db.flush()
            (await db.get(NoteContentTemplate, ids["tpl"])).current_version_id = newer.id
            await db.commit()
        assert (await _note(env, "basic")).text_content == v1_body  # 发布新版本不碰已引用项目

        v2_source = {"source_kind": "content_template", "template_version_id": str(newer.id),
                     "template_content_hash": "c" * 64}
        v2_preview = (await _preview(env, v2_source)).json()
        v2_item = next(i for i in v2_preview["items"] if i["target_note_section"] == "一、1")
        assert (await _note(env, "basic")).text_content == v1_body  # 只预览不写
        resp = await _apply(env, v2_source, v2_preview, [_selection(v2_item)], mutation_id="mut-v2-000001")
        assert resp.status_code == 200, resp.text
        assert (await _note(env, "basic")).text_content == "<p>新版本正文</p>"
        assert {uuid.UUID(str(r.template_version_id)) for r in await _applications(env)} == {ids["ver"], newer.id}

    @pytest.mark.asyncio
    async def test_mutation_id_required(self, env):
        source, preview, item = await _previewed(env)
        resp = await _apply(env, source, preview, [_selection(item)], mutation_id=None)
        _assert_error(resp, 422, "MUTATION_ID_REQUIRED")
        assert await _applications(env) == []

    @pytest.mark.asyncio
    async def test_target_changed_after_preview_is_409_with_sections(self, env):
        source, preview, item = await _previewed(env)
        async with env["factory"]() as db:
            (await db.get(DisclosureNote, env["ids"]["basic"])).text_content = "<p>他人刚改过</p>"
            await db.commit()
        resp = await _apply(env, source, preview, [_selection(item)])
        assert _assert_error(resp, 409, "NOTE_BATCH_STALE_PREVIEW")["stale_notes"] == ["一、1"]
        assert (await _note(env, "basic")).text_content == "<p>他人刚改过</p>"
        assert await _applications(env) == [] and env["published"] == []

    @pytest.mark.asyncio
    async def test_source_changed_after_preview_is_409(self, env):
        source, preview, item = await _previewed(env)
        other = next(s for s in source["parsed_source"]["sections"] if s["title"] == "会计期间")
        other["html"] = "<p>预览后被改过的来源</p>"
        resp = await _apply(env, source, preview, [_selection(item)])
        _assert_error(resp, 409, "PREVIEW_DIGEST_MISMATCH")
        assert (await _note(env, "basic")).text_content == "<p>旧基本情况</p>"

    @pytest.mark.parametrize("gate,status", [("readonly", 403), ("consol_lock", 423), ("archived", 423)])
    @pytest.mark.asyncio
    async def test_write_gates(self, env, gate, status):
        # 知识库来源不绑定解析人：去掉门禁后 readonly 用户会真写进去（上传来源的签名绑定用户，会先 422）
        source, preview, item = await _previewed(env, "knowledge")
        if gate == "readonly":
            env["act_as"]("viewer")
        elif gate == "consol_lock":
            await _update_project(env, consol_lock=True)
        else:
            await _update_project(env, status=ProjectStatus.archived)
        resp = await _apply(env, source, preview, [_selection(item)])
        assert resp.status_code == status, resp.text
        assert (await _note(env, "basic")).text_content == "<p>旧基本情况</p>"


class TestTextSources:
    """Task 50 / AC-11.3：编辑器来源标记 —— 各章最近一次批量写入的来源（只读，不进正文）。"""

    @pytest.mark.asyncio
    async def test_latest_source_per_note_with_writer_and_modified_flag(self, env):
        ids, url = env["ids"], _url(env, "/disclosure-note-text-batch/sources")
        async with env["factory"]() as db:
            db.add(StaffMember(name="审计助理甲", user_id=ids["me"]))
            await db.commit()
        source, preview, item = await _previewed(env)
        assert (await _apply(env, source, preview, [_selection(item)])).status_code == 200
        async with env["factory"]() as db:  # 写入记录与章节同一时刻（取在加锁后），不用库默认的事务开始时间
            (application,) = (await db.execute(sa.select(NoteTextBatchApplication))).scalars().all()
            note = await db.get(DisclosureNote, uuid.UUID(item["target_note_id"]))
            assert application.applied_at == note.updated_at
        items = (await env["client"].get(url, params={"year": YEAR})).json()["items"]
        assert set(items) == {item["target_note_id"]}  # 只列被批量写入过的章节
        entry = items[item["target_note_id"]]
        assert (entry["source_kind"], entry["file_name"], entry["mode"], entry["applied_by_name"],
                entry["modified_after"]) == ("upload", "附注.docx", "replace", "审计助理甲", False)

        async with env["factory"]() as db:  # 写入后有人改了正文 ⇒ 标「之后有修改」
            (await db.get(DisclosureNote, uuid.UUID(item["target_note_id"]))).text_content = "<p>人工改过</p>"
            await db.commit()
        assert (await env["client"].get(url, params={"year": YEAR})).json()["items"][
            item["target_note_id"]]["modified_after"] is True
        assert (await env["client"].get(url, params={"year": YEAR + 1})).json()["items"] == {}

    @pytest.mark.asyncio
    async def test_template_source_carries_name_and_version_and_latest_wins(self, env):
        upload, preview, item = await _previewed(env)
        assert (await _apply(env, upload, preview, [_selection(item)], mutation_id="mut-src-1")).status_code == 200
        tpl = {"source_kind": "content_template", "template_version_id": str(env["ids"]["ver"]),
               "template_content_hash": TPL_HASH}
        tpl_preview = (await _preview(env, tpl)).json()
        tpl_item = next(i for i in tpl_preview["items"] if i["target_note_id"] == item["target_note_id"])
        assert (await _apply(env, tpl, tpl_preview, [_selection(tpl_item)], mutation_id="mut-src-2")).status_code == 200
        entry = (await env["client"].get(_url(env, "/disclosure-note-text-batch/sources"),
                                         params={"year": YEAR})).json()["items"][item["target_note_id"]]
        assert (entry["source_kind"], entry["template_name"], entry["version_no"]) == ("content_template", "乙公司全套", 1)

    @pytest.mark.asyncio
    async def test_newest_by_time_knowledge_shape_and_hidden_rows(self, env):
        """「最近一次」按写入时间而非 id；知识库只回文档名与版本；已删章节 / 他项目章节 / 坏行不展示也不报错。"""
        ids = env["ids"]

        def application(app_id: str, kind: str, meta: dict, at: datetime, detail: list) -> NoteTextBatchApplication:
            return NoteTextBatchApplication(
                id=uuid.UUID(app_id), source_kind=kind, source_metadata=meta, target_project_id=ids["p"],
                target_year=YEAR, mutation_id=f"mut-{app_id[:2]}", request_fingerprint="f" * 64,
                preview_digest="d" * 64, applied_count=len(detail), detail=detail, applied_by=ids["me"], applied_at=at)

        basic_hash = hashlib.sha256("<p>旧基本情况</p>".encode()).hexdigest()
        async with env["factory"]() as db:
            (await db.get(DisclosureNote, ids["cash"])).is_deleted = True
            db.add_all([  # id 大的反而更早：按 id 排就会选错
                application("ffffffff-ffff-4fff-bfff-ffffffffffff", "upload", {"file_name": "旧.docx"},
                            datetime(2026, 1, 1, tzinfo=timezone.utc),
                            [{"target_note_id": str(ids["basic"]), "mode": "replace", "after_hash": basic_hash}]),
                application("00000000-0000-4000-8000-000000000001", "knowledge",
                            {"documents": [{"doc_id": str(ids["doc"]), "document_name": "政策汇编.md", "doc_version": 3}]},
                            datetime(2026, 2, 1, tzinfo=timezone.utc),
                            [{"target_note_id": str(ids["basic"]), "mode": "append", "after_hash": basic_hash},
                             {"target_note_id": str(ids["cash"]), "mode": "append", "after_hash": "x"},
                             {"target_note_id": str(ids["foreign"]), "mode": "append", "after_hash": "x"},
                             "坏行", {"mode": "append"}, {"target_note_id": "不是-uuid"}, {"target_note_id": None}]),
            ])
            await db.commit()
        resp = await env["client"].get(_url(env, "/disclosure-note-text-batch/sources"), params={"year": YEAR})
        assert resp.status_code == 200, resp.text
        items = resp.json()["items"]
        assert set(items) == {str(ids["basic"])}
        entry = items[str(ids["basic"])]
        assert (entry["source_kind"], entry["mode"], entry["modified_after"]) == ("knowledge", "append", False)
        assert entry["documents"] == [{"document_name": "政策汇编.md", "doc_version": 3}]
        assert entry["application_id"] == "00000000-0000-4000-8000-000000000001"

    @pytest.mark.asyncio
    async def test_readonly_member_reads_outsider_forbidden(self, env):
        url = _url(env, "/disclosure-note-text-batch/sources")
        env["act_as"]("viewer")
        assert (await env["client"].get(url, params={"year": YEAR})).status_code == 200
        env["act_as"]("outsider")
        assert (await env["client"].get(url, params={"year": YEAR})).status_code == 403

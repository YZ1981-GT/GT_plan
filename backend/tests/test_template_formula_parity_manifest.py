# -*- coding: utf-8 -*-
"""前端 ↔ 模板公式对拍清单守卫（spec oo-writeback-template-parity-closure R2.4，design §3.3）。

清单 ``templateFormulaParity.generated.json`` 由 ``scripts/gen/gen_template_formula_parity.py`` 从生产契约 +
权威模板生成，前端 ``templateFormulaParity.spec.ts`` 用它与前端行计算逐值对拍。本守卫保证：

① 入库清单 = 生成器现算（等价 ``--check``）：契约或模板改了却没重生成就红；
② 换算口径：同行引用换成字段名；他行 / 他表 / 不支持的函数标 ``skip`` 并写原因（不猜）；
   非受管公式列在同行展开、表头也空的间隔列按 0；
③ 变异：模板公式改一格，生成结果随之改变（``--check`` 红）。
"""
from __future__ import annotations

import importlib.util
import json
import shutil
import sys
from pathlib import Path

import openpyxl
import pytest

BACKEND = Path(__file__).resolve().parents[1]
_GEN = BACKEND / "scripts" / "gen" / "gen_template_formula_parity.py"
_OUT = BACKEND.parent / "audit-platform/frontend/src/components/workpaper/sync/__tests__/templateFormulaParity.generated.json"


def _load_gen():
    name = "gen_template_formula_parity_t"
    spec = importlib.util.spec_from_file_location(name, _GEN)
    module = importlib.util.module_from_spec(spec)
    # 生成器里有 @dataclass：dataclasses 按 cls.__module__ 去 sys.modules 取模块，未登记会抛 AttributeError
    sys.modules[name] = module
    spec.loader.exec_module(module)  # type: ignore[union-attr]
    return module


@pytest.fixture(scope="module")
def gen():
    return _load_gen()


@pytest.fixture(scope="module")
def payload(gen):
    return gen.collect(BACKEND)


def test_manifest_is_generator_output(gen, payload):
    """入库清单逐字等于生成器现算结果（等价 ``--check``）。"""
    assert _OUT.read_text(encoding="utf-8") == gen.render(payload), (
        "清单与契约 / 模板现算不一致：在 backend 目录跑 scripts/gen/gen_template_formula_parity.py 重生成"
    )


def test_check_mode_exit_codes(gen, payload, tmp_path, capsys):
    """``--check``：一致 exit 0 并报 [ok]；改一个字节 exit 2 并报 [DRIFT]（输出内容一并断言，不只看退出码）。"""
    same = tmp_path / "same.json"
    same.write_text(gen.render(payload), encoding="utf-8", newline="\n")
    assert gen.main(["--check", "--out", str(same)]) == 0
    assert "[ok]" in capsys.readouterr().out

    drift = tmp_path / "drift.json"
    drift.write_text(gen.render(payload).replace('"ok": ', '"ok": 1', 1), encoding="utf-8", newline="\n")
    assert gen.main(["--check", "--out", str(drift)]) == 2
    assert "[DRIFT]" in capsys.readouterr().out


def test_manifest_is_self_consistent(payload):
    """每个公式字段要么有表达式、要么写明跳过原因；计数与条目一致；条目的 store item 都登记在 items 里。"""
    entries = payload["entries"]
    assert entries, "清单为空：契约目录或模板路径不对"
    bad = [e for e in entries if (e["expr"] is None) == (e["skip"] is None)]
    assert bad == []
    ok = sum(1 for e in entries if e["expr"] is not None)
    assert payload["summary"] == {"formula_fields": len(entries), "ok": ok, "skip": len(entries) - ok}
    assert {e["store_item"] for e in entries} <= set(payload["items"])
    for item in payload["items"].values():
        overlap = set(item["inputs"]) & set(item["auto_fields"])
        assert not overlap, f"auto_source 字段不能同时当录入：{overlap}"


def _ctx(gen, ws, row: int, columns: dict[str, str]):
    ctx = gen.TableCtx("c", ws.title, "item", "t", row)
    for col, field in columns.items():
        ctx.columns[col] = gen.Column(field, "editable", "amount", "")
    ctx.ws = ws
    return ctx


def test_convert_same_row_refs_become_field_names(gen):
    wb = openpyxl.Workbook()
    ws = wb.active
    ctx = _ctx(gen, ws, 5, {"B": "opening", "C": "inc", "D": "dec", "E": "flag"})
    assert gen.convert("=B5+SUM(C5:D5)-D5", ctx) == "opening + SUM( inc, dec ) - dec"
    assert gen.convert('=IF(E5="是",B5,0)', ctx) == 'IF( flag = "是" , opening , 0 )'
    assert "是" in ctx.literals


@pytest.mark.parametrize(
    ("formula", "reason"),
    [
        ("=B5+B6", "引用他行"),
        ("=B5+$B$20", "引用他行"),
        ("='明细表'!B5+C5", "引用他表"),
        ("=SUMPRODUCT(B5:C5)", "不在支持范围"),
        ('=B5&"元"', "字符串拼接"),
    ],
)
def test_convert_skips_what_it_cannot_express(gen, formula, reason):
    """换不成同行字段表达式的一律跳过并写原因，不猜。"""
    wb = openpyxl.Workbook()
    ws = wb.active
    ctx = _ctx(gen, ws, 5, {"B": "opening", "C": "inc"})
    with pytest.raises(gen.Skip, match=reason):
        gen.convert(formula, ctx)


def test_convert_divides_percent_points_columns_by_100(gen):
    """百分点列（store 存 3.6 = 3.6%，模板格存 0.036）：公式里引用它时换回小数（F3-4 应计利息 = ROUND(F*G,2)）；
    百分点列也算数值录入（随机取数，按录入核对）。"""
    wb = openpyxl.Workbook()
    ws = wb.active
    ctx = _ctx(gen, ws, 10, {"F": "faceValue"})
    ctx.columns["G"] = gen.Column("interestRate", "editable", "percent_points", "票面利率")
    assert gen.convert("=ROUND(F10*G10,2)", ctx) == "ROUND( faceValue * ( interestRate / 100 ) , 2 )"
    assert "percent_points" in gen.NUMERIC_TYPES


def test_convert_inlines_unmanaged_formula_columns_and_zeroes_blank_gaps(gen):
    """非受管的公式列（小计列）同行展开；数据格与表头都空的间隔列按 0；有表头的非受管录入列跳过。"""
    wb = openpyxl.Workbook()
    ws = wb.active
    ws["D5"] = "=B5+C5"
    ws["G4"] = "其他"
    ctx = _ctx(gen, ws, 5, {"B": "a", "C": "b", "E": "c"})
    ctx.blank_columns.add("F")
    assert gen.convert("=D5-E5+F5", ctx) == "( a + b ) - c + 0"
    with pytest.raises(gen.Skip, match="引用非受管列 G"):
        gen.convert("=E5+G5", ctx)


def test_template_formula_change_changes_manifest(gen, tmp_path):
    """变异：D5-2 模板「期末审定」O 列 +RJE 改成 −RJE ⇒ 生成的表达式随之改变。"""
    contract_path = BACKEND / "data" / "workpaper_sync_contracts" / "d5.receivables_financing_detail.json"
    contract = json.loads(contract_path.read_text(encoding="utf-8"))
    rel = contract["template"]["relative_path"]
    root = tmp_path / "backend"
    (root / "data" / "workpaper_sync_contracts").mkdir(parents=True)
    shutil.copy(contract_path, root / "data" / "workpaper_sync_contracts" / contract_path.name)
    (root / "wp_templates" / rel).parent.mkdir(parents=True)
    shutil.copy(BACKEND / "wp_templates" / rel, root / "wp_templates" / rel)

    def end_audited_expr() -> str:
        entries = gen.collect(root)["entries"]
        hit = [e for e in entries if e["store_item"] == "D5-2-rows" and e["field"] == "endAudited"]
        assert len(hit) == 1
        return hit[0]["expr"]

    assert end_audited_expr() == "endUnadjusted + endAje + endRje"
    wb = openpyxl.load_workbook(root / "wp_templates" / rel)
    ws = wb["应收款项融资明细表D5-2"]
    assert ws["O12"].value == "=L12+M12+N12"
    ws["O12"] = "=L12+M12-N12"
    wb.save(root / "wp_templates" / rel)
    assert end_audited_expr() == "endUnadjusted + endAje - endRje"

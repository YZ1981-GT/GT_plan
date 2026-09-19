"""两个模板入口的必填列兼容性 —— adjustment-import-export-contract / Task 1.5（Requirement 4）。

背景：调整分录有两个模板下载入口
  - 工具栏 / 导入弹窗（本任务改指向）：`GET /api/projects/{pid}/adjustments/export-template`
    项目感知富模板（4 sheet：关注事项 / AJE模板 / RJE模板 / 项目科目库，E 列科目下拉 + D/F/G/H 联动）
  - 无项目上下文降级：`GET /api/import-templates/adjustments/download`
    通用裸模板（`generate_template` 按 `TEMPLATE_COLUMNS` 产出，sheet「数据」+「填写说明」，无下拉）

R4.3 要求「任一入口下载的模板填写并导入 THEN 系统 SHALL 均能接受（两种模板的必填列集合保持互相兼容）」。
本文件把该兼容性锁成契约测试：

  P4-1  两入口列头集合逐字一致（裸模板表头带 `* ` 必填前缀，去前缀后比较）
  P4-2  `TEMPLATE_COLUMNS` 声明的必填列在两入口表头里都存在
  P4-3  富模板填一行真实数据 → `validate_import_file` 通过且 `parse_import_data` 读到该行
  P4-4  裸模板填一行真实数据 → 同上
  P4-5  两入口同一笔业务数据解析出的业务字段逐字等价（类型/摘要/科目名称/金额不失真）

不触库：富模板复用 characterization 里的 fake DB helper（`_build_rich_template`），
裸模板走纯函数 `generate_template`。
"""
from __future__ import annotations

from io import BytesIO

import openpyxl
import pytest

from app.services.import_template_service import (
    TEMPLATE_COLUMNS,
    ImportType,
    generate_template,
    parse_import_data,
    validate_import_file,
)
from tests.test_adjustment_ie_characterization import _build_rich_template

_ADJ_COLUMNS = TEMPLATE_COLUMNS[ImportType.adjustments]
_EXPECTED_NAMES = [c[0] for c in _ADJ_COLUMNS]
_REQUIRED_NAMES = {c[0] for c in _ADJ_COLUMNS if c[1]}

# 测试数据行：用唯一摘要作定位标记，避免与模板内置示例混淆
_MARK = "Task1.5 双入口兼容性测试分录"
_ROW = {
    "编号": "AJE-901",
    "类型": "AJE",
    "摘要": _MARK,
    "二级科目名称": "应收账款",
    "借方金额": 12345.67,
}


def _header_texts(ws) -> list[str]:
    """读第 1 行列头（去掉裸模板的 `* ` 必填前缀，与读取端 lstrip 口径一致）。"""
    out: list[str] = []
    for cell in next(ws.iter_rows(min_row=1, max_row=1, values_only=False), []):
        out.append(str(cell.value or "").strip().lstrip("* ").strip())
    while out and out[-1] == "":
        out.pop()
    return out


def _fill_row(ws, row_idx: int, headers: list[str]) -> None:
    """按列头把 `_ROW` 写到指定行（未登记的列不动，保留模板公式/空白）。"""
    for col_idx, name in enumerate(headers, 1):
        if name in _ROW:
            ws.cell(row=row_idx, column=col_idx, value=_ROW[name])


def _save(wb) -> bytes:
    buf = BytesIO()
    wb.save(buf)
    return buf.getvalue()


def _find_marked(rows: list[dict]) -> dict | None:
    for r in rows:
        if str(r.get("摘要") or "").strip() == _MARK:
            return r
    return None


def _bare_template_workbook():
    return openpyxl.load_workbook(BytesIO(generate_template(ImportType.adjustments)))


# ═══════════════════════════════════════════════════════════════════════════
# P4-1 / P4-2 列头与必填列集合
# ═══════════════════════════════════════════════════════════════════════════


def test_p4_2_bare_template_contains_all_required_columns():
    """裸模板（generate_template）表头 = TEMPLATE_COLUMNS 全列，必填列齐备。"""
    wb = _bare_template_workbook()
    headers = _header_texts(wb["数据"])
    assert headers == _EXPECTED_NAMES
    assert _REQUIRED_NAMES <= set(headers)


@pytest.mark.parametrize("sheet", ["AJE模板", "RJE模板"])
async def test_p4_1_rich_and_bare_headers_are_identical(monkeypatch, sheet):
    """富模板录入 sheet 列头与裸模板逐字一致 → 必填列集合天然互相兼容。"""
    rich = await _build_rich_template(monkeypatch)
    bare_headers = _header_texts(_bare_template_workbook()["数据"])
    rich_headers = _header_texts(rich[sheet])

    assert rich_headers == bare_headers, f"{sheet} 与裸模板列头不一致 → 两入口不兼容"
    assert _REQUIRED_NAMES <= set(rich_headers)


# ═══════════════════════════════════════════════════════════════════════════
# P4-3 / P4-4 两入口填写后都能被导入接受
# ═══════════════════════════════════════════════════════════════════════════


async def test_p4_3_rich_template_filled_is_accepted(monkeypatch):
    """富模板 AJE模板 sheet 填一行真实数据（示例区之后）→ 校验通过 + 解析读到该行。"""
    wb = await _build_rich_template(monkeypatch)
    ws = wb["AJE模板"]
    _fill_row(ws, 7, _header_texts(ws))  # 2~6 为示例/空白区，7 起为用户数据
    data = _save(wb)

    result = validate_import_file(ImportType.adjustments, data, "rich.xlsx")
    assert result.valid, f"富模板填写后校验未通过: {[e.message for e in result.errors]}"

    rows = parse_import_data(ImportType.adjustments, data)
    hit = _find_marked(rows)
    assert hit is not None, "富模板填写的数据行未被解析出来"
    assert str(hit.get("类型")).strip().upper() == "AJE"


def test_p4_4_bare_template_filled_is_accepted():
    """裸模板「数据」sheet 填一行真实数据 → 校验通过 + 解析读到该行（含类型列）。"""
    wb = _bare_template_workbook()
    ws = wb["数据"]
    _fill_row(ws, 3, _header_texts(ws))  # 第 2 行为内置示例，第 3 行起可填
    data = _save(wb)

    result = validate_import_file(ImportType.adjustments, data, "bare.xlsx")
    assert result.valid, f"裸模板填写后校验未通过: {[e.message for e in result.errors]}"

    rows = parse_import_data(ImportType.adjustments, data)
    hit = _find_marked(rows)
    assert hit is not None, "裸模板填写的数据行未被解析出来"
    assert str(hit.get("类型")).strip().upper() == "AJE"


# ═══════════════════════════════════════════════════════════════════════════
# P4-5 同一笔业务数据在两入口解析等价
# ═══════════════════════════════════════════════════════════════════════════


async def test_p4_5_both_entries_parse_to_equivalent_business_row(monkeypatch):
    """同一笔分录经两入口模板导入后，业务字段（类型/摘要/科目名称/借方金额）逐字等价。"""
    rich = await _build_rich_template(monkeypatch)
    rich_ws = rich["AJE模板"]
    _fill_row(rich_ws, 7, _header_texts(rich_ws))
    rich_hit = _find_marked(parse_import_data(ImportType.adjustments, _save(rich)))

    bare = _bare_template_workbook()
    bare_ws = bare["数据"]
    _fill_row(bare_ws, 3, _header_texts(bare_ws))
    bare_hit = _find_marked(parse_import_data(ImportType.adjustments, _save(bare)))

    assert rich_hit is not None and bare_hit is not None

    def _business(row: dict) -> dict:
        return {
            "编号": str(row.get("编号") or "").strip(),
            "类型": str(row.get("类型") or "").strip().upper(),
            "摘要": str(row.get("摘要") or "").strip(),
            "二级科目名称": str(row.get("二级科目名称") or "").strip(),
            "借方金额": float(row.get("借方金额") or 0),
            "贷方金额": float(row.get("贷方金额") or 0),
        }

    assert _business(rich_hit) == _business(bare_hit)

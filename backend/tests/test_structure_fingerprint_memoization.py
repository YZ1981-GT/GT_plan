# -*- coding: utf-8 -*-
"""``structure_fingerprint`` 按字节摘要记忆化的守卫。

2026-09-22 冷路径实测：`attach_pilot_*` / `register_from_manifest` →
`observe_published_frozen_definitions` → `collect_workbook_structure` → 本函数，
同一份不变字节被反复完整解析（cProfile 里 `openpyxl.load_workbook` 累计 60 秒）。

记忆化只有在三条同时成立时才是安全的，缺一条都会从优化变成数据正确性问题：

1. **同字节只算一次**（否则优化没生效）；
2. **不同字节不串**（key 是 `sha256(data)`，串了就等于读到另一个底稿的结构）；
3. **交出去的对象被就地改动不污染缓存** —— `WorkbookFingerprint` 是可变 dataclass
   （内部全是 dict/list），把缓存对象原样交出去时任何调用方的就地修改都会毒化后续全部命中。
"""
from __future__ import annotations

import io
import zipfile

import openpyxl
import pytest

from app.services import excel_structure_fingerprint as MOD
from app.services.excel_structure_fingerprint import (
    FingerprintError,
    clear_structure_fingerprint_cache,
    structure_fingerprint,
)


@pytest.fixture(autouse=True)
def _clean_cache():
    clear_structure_fingerprint_cache()
    yield
    clear_structure_fingerprint_cache()


@pytest.fixture()
def parse_counter(monkeypatch: pytest.MonkeyPatch) -> list[int]:
    """记录真实解析次数（缓存未命中才会 +1）。"""
    calls: list[int] = []
    real = MOD._structure_fingerprint_uncached

    def counting(data: bytes):  # type: ignore[no-untyped-def]
        calls.append(len(data))
        return real(data)

    monkeypatch.setattr(MOD, "_structure_fingerprint_uncached", counting)
    return calls


def _workbook_bytes(*, a1: str = "hello", sheet: str = "S1") -> bytes:
    wb = openpyxl.Workbook()
    ws = wb.active
    ws.title = sheet
    ws["A1"] = a1
    buf = io.BytesIO()
    wb.save(buf)
    wb.close()
    return buf.getvalue()


# ═══ 1. 同字节只算一次 ═══


def test_identical_bytes_are_parsed_only_once(parse_counter) -> None:
    data = _workbook_bytes()

    first = structure_fingerprint(data)
    for _ in range(5):
        structure_fingerprint(data)

    assert len(parse_counter) == 1, (
        f"同一份字节解析了 {len(parse_counter)} 次 —— 记忆化未生效"
    )
    assert first.errors == []


def test_equal_but_distinct_byte_objects_still_hit_the_cache(parse_counter) -> None:
    """key 是内容摘要而非对象身份：复制一份 bytes 仍必须命中。"""
    data = _workbook_bytes()
    structure_fingerprint(data)
    structure_fingerprint(bytes(bytearray(data)))
    assert len(parse_counter) == 1


# ═══ 2. 不同字节不串 ═══


def test_different_bytes_do_not_share_a_cache_entry(parse_counter) -> None:
    """🔴 两份字节必须**等长**。

    长度不同的一对没有区分力：把 key 从 `sha256(data)` 换成 `len(data)` 这类退化实现
    照样能让它通过（本轮变异检验实测 SURVIVED 过一次）。等长不同内容才真正测「key 是
    内容摘要」。
    """
    one = _workbook_bytes(a1="AAA")
    two = _workbook_bytes(a1="BBB")
    assert len(one) == len(two) and one != two, (
        "夹具失效：需要一对等长但字节不同的 workbook 才能测出 key 退化"
    )

    got_one = structure_fingerprint(one)
    got_two = structure_fingerprint(two)

    assert len(parse_counter) == 2, f"不同字节应各算一次：{parse_counter}"
    assert got_one.part_digests != got_two.part_digests, (
        "两份不同 workbook 拿到了相同的部件摘要 —— 缓存串了，会读到另一个底稿的结构"
    )


def test_cache_eviction_is_bounded() -> None:
    """LRU 有上限：连续喂超过上限份字节不得无界增长。"""
    limit = MOD._FINGERPRINT_CACHE_MAX
    for index in range(limit + 4):
        structure_fingerprint(_workbook_bytes(a1=f"value-{index}"))
    assert len(MOD._FINGERPRINT_CACHE) <= limit, (
        f"缓存涨到 {len(MOD._FINGERPRINT_CACHE)} 份，超过上限 {limit}"
    )


# ═══ 3. 返回值被就地修改不得污染缓存 ═══


def test_mutating_a_returned_fingerprint_does_not_poison_later_hits() -> None:
    data = _workbook_bytes()

    first = structure_fingerprint(data)
    baseline_parts = dict(first.part_digests)
    first.part_digests["xl/INJECTED.xml"] = "tampered"
    first.part_digests.pop(next(iter(baseline_parts)), None)

    second = structure_fingerprint(data)
    assert "xl/INJECTED.xml" not in second.part_digests, (
        "调用方对返回值的就地修改漏进了缓存 —— 后续每次命中都会拿到被篡改的结构事实"
    )
    assert second.part_digests == baseline_parts, (
        "命中值与首次计算值不再相等 —— 缓存已被上一次调用毒化"
    )


def test_two_callers_get_independent_objects() -> None:
    data = _workbook_bytes()
    a = structure_fingerprint(data)
    b = structure_fingerprint(data)
    assert a is not b, "两次调用返回同一个对象 —— 一方改动会影响另一方"
    assert a.part_digests == b.part_digests


def test_mutating_a_cache_hit_result_does_not_poison_the_next_hit() -> None:
    """🔴 存入侧与命中侧是**两处** deepcopy，判据必须分别覆盖。

    只改「首次返回值」测不到命中侧：首次返回的是刚算出来的 `computed`，存入时已经拷过一份，
    所以即使命中侧把缓存对象原样交出去，那条判据照样绿（本轮变异检验实测 SURVIVED 过一次）。
    这里改的是**第二次**（命中）拿到的对象，再看第三次。
    """
    data = _workbook_bytes()

    structure_fingerprint(data)  # 首次：算并入缓存
    hit = structure_fingerprint(data)  # 第二次：命中
    baseline_parts = dict(hit.part_digests)
    hit.part_digests["xl/INJECTED_ON_HIT.xml"] = "tampered"

    third = structure_fingerprint(data)
    assert "xl/INJECTED_ON_HIT.xml" not in third.part_digests, (
        "对**命中值**的就地修改漏进了缓存 —— 命中侧缺 deepcopy，"
        "后续每次命中都拿到被篡改的结构事实"
    )
    assert third.part_digests == baseline_parts


# ═══ 4. 失败路径不得进缓存 ═══


def test_invalid_bytes_still_raise_and_are_not_cached(parse_counter) -> None:
    junk = b"not-a-zip-at-all"
    for _ in range(3):
        with pytest.raises(FingerprintError):
            structure_fingerprint(junk)
    assert len(parse_counter) == 3, (
        "失败结果被缓存了 —— 抛错路径不得占用 key（否则同一字节以后连错误类型都可能漂移）"
    )


def test_corrupt_zip_is_not_silently_turned_into_an_empty_fingerprint() -> None:
    buf = io.BytesIO()
    with zipfile.ZipFile(buf, "w") as zf:
        zf.writestr("random.txt", "not a workbook")
    with pytest.raises(FingerprintError):
        structure_fingerprint(buf.getvalue())

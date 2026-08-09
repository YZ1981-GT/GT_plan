"""Wave 5 守卫：自定义底稿独立导出路径。

spec: .kiro/specs/custom-workpaper-dual-mode-formula-and-batch/
  Task 16（独立导出服务）/ Task 17（两条路径分流）/ Task 18（本守卫）

覆盖 Property：
  10 custom 导出不走 load_schema / export_workpaper_xlsx；分流早于 load_schema
  18 两条导出路径都分流（`/export-xlsx` 与 `WpExportEngine.export_single`）
  19 格式判定显式登记 `"custom"`
  20 往返可逆（导出字节流可被 openpyxl 打开，且 cells 逐格相等）

🔴 判据设计要点（memory 铁律）：
  - **必须先 stripComments** —— 生产代码注释里逐字引用了
    `except (TemplateNotFoundError, Exception)` 用于解释分流理由；不剥注释会让
    「分流位置」判据把注释行当成真实回退分支（实测行 473 是注释、行 493 才是真分支）。
  - 位置类判据比**调用点**（`name(`）而非裸标识符 —— 函数体内的局部 import
    也含该标识符（本 spec 已两次踩中）。
  - 反向自检必须证明禁令有区分度（标准底稿路径**仍然**引用 load_schema）。
"""

from __future__ import annotations

import re
from io import BytesIO
from pathlib import Path
from typing import Any

import pytest

from app.services.custom_workpaper_export import (
    CustomExportError,
    export_custom_workpaper,
)

_BACKEND = Path(__file__).resolve().parents[1]
_REPO = _BACKEND.parent


def _strip_py_comments(src: str) -> str:
    """剥 docstring 与 `#` 注释。

    🔴 本 spec 的生产代码注释里逐字引用了被禁/被解释的符号
    （`except (TemplateNotFoundError, Exception)` / `load_schema`），
    不剥会让守卫把说明文字当成真实代码 —— 位置判据会算错、禁令判据会假红。
    """
    src = re.sub(r'"""[\s\S]*?"""', '""', src)
    src = re.sub(r"'''[\s\S]*?'''", "''", src)
    src = re.sub(r"(?m)#.*$", "", src)
    return src


def _read(rel: str) -> str:
    p = _REPO / rel
    assert p.exists(), f"守卫判据文件缺失: {rel}"
    txt = p.read_text(encoding="utf-8")
    assert txt.strip(), f"文件为空（判据失效）: {rel}"
    return txt


def _fn_body(src: str, name: str) -> str:
    """按缩进截取 python 函数体（缩进用 `[ \\t]*`，禁 `\\s*`）。"""
    m = re.search(rf"(?m)^([ \t]*)(?:async\s+)?def\s+{re.escape(name)}\s*\(", src)
    assert m, f"未找到函数 {name}（判据失效）"
    indent = len(m.group(1))
    lines = src[m.start():].splitlines()
    out = [lines[0]]
    for ln in lines[1:]:
        if ln.strip() and (len(ln) - len(ln.lstrip())) <= indent and re.match(
            r"\s*(@|(?:async\s+)?def |class )", ln
        ):
            break
        out.append(ln)
    body = "\n".join(out)
    assert len(body) > 50, f"{name} 函数体过短（截取失败）"
    return body


def _fallback_except_pos(body: str) -> int:
    """定位「空白 workbook 回退」那条 `except` 的位置。

    🔴 为什么按**形态**而非写死字面量（2026-08-09 假红 + 假绿双证）：
      改造前判据写死 `except (TemplateNotFoundError, Exception)`，而本轮把它
      收窄成 `except Exception:  # noqa: BLE001`（元组里带 `Exception` 等于吞
      一切、且 flake8 不告警 = 更坏）。字面量判据于是：
        · `test_engine_gate_precedes_...` 直接**假红** —— 看着像「回退分支被删了」；
        · `test_engine_custom_branch_not_wrapped_in_try` **假绿** ——
          `find()` 返回 `-1`，`body[gate:-1]` 恰好仍是「分流到末尾」这段，
          断言照样通过 ⇒ 同一处过期判据在两个测试里表现相反。
      形态匹配同时认两种写法，且返回 -1 时由调用方显式断言，不再靠切片兜。
    """
    m = re.search(r"(?m)^\s*except\s+(?:Exception|\(\s*[\w.]+\s*,\s*Exception\s*\))\s*(?:as\s+\w+\s*)?:", body)
    return m.start() if m else -1


def _call_pos(body: str, name: str) -> int:
    """调用点位置（-1 = 无调用）。裸标识符会命中局部 import，必须匹配 `name(`。"""
    m = re.search(rf"{re.escape(name)}\s*\(", body)
    return m.start() if m else -1


class _FakeWp:
    def __init__(self, file_path: str | None = None):
        self.file_path = file_path
        self.file_version = 1


def _make_xlsx(tmp_path: Path, sheet: str, cells: dict[str, Any]) -> Path:
    import openpyxl

    wb = openpyxl.Workbook()
    ws = wb.active
    ws.title = sheet
    for ref, val in cells.items():
        ws[ref] = val
    p = tmp_path / "custom.xlsx"
    wb.save(str(p))
    wb.close()
    return p


_ROUTER_REL = "backend/app/routers/wp_xlsx_export.py"
_ENGINE_REL = "backend/app/services/wp_export/export_engine.py"
_SERVICE_REL = "backend/app/services/custom_workpaper_export.py"


# ════════════════════════════════════════════════════════════════════════════
class TestGuardSelfCheck:
    def test_strip_comments_removes_quoted_symbols(self):
        """🔴 证明剥注释确实生效 —— 否则位置判据会把注释行当真实分支。"""
        s = 'def f():\n    # except (TemplateNotFoundError, Exception) 说明\n    x = 1'
        c = _strip_py_comments(s)
        assert "TemplateNotFoundError" not in c

    def test_call_pos_ignores_local_import(self):
        """🔴 裸标识符会命中函数体内的局部 import（本 spec 已两次踩中）。"""
        body = "def f():\n    from m import write_cells_to_xlsx\n    return 1"
        assert _call_pos(body, "write_cells_to_xlsx") == -1
        body2 = body + "\n    write_cells_to_xlsx(a, b, c)"
        assert _call_pos(body2, "write_cells_to_xlsx") > 0

    def test_read_missing_raises(self):
        with pytest.raises(AssertionError):
            _read("backend/__no_such_file__.py")


# ════════════════════════════════════════════════════════════════════════════
class TestProperty10CustomExportAvoidsSchema:
    """Property 10：custom 导出不依赖 render schema。"""

    def test_service_does_not_reference_schema_or_legacy_exporter(self):
        src = _strip_py_comments(_read(_SERVICE_REL))
        assert "load_schema" not in src, "custom 导出服务不得引用 load_schema"
        assert "export_workpaper_xlsx" not in src, (
            "custom 导出服务不得走 dynamic_table 导出器（它读 rows，custom 是 cells）"
        )

    def test_reverse_selfcheck_standard_path_still_uses_load_schema(self):
        """🔴 反向自检：标准底稿路径**仍然**引用 load_schema（证明禁令有区分度）。"""
        router = _strip_py_comments(_read(_ROUTER_REL))
        assert "load_schema" in router, (
            "标准路径应仍用 load_schema —— 否则上面那条禁令没有区分度"
        )

    def test_router_gate_precedes_load_schema(self):
        """分流判定必须在 load_schema 之前，否则仍会 500。"""
        src = _strip_py_comments(_read(_ROUTER_REL))
        body = _fn_body(src, "export_xlsx")
        gate = _call_pos(body, "resolve_is_custom")
        schema = _call_pos(body, "load_schema")
        assert gate >= 0, "export_xlsx 应有 custom 分流"
        assert schema >= 0, "export_xlsx 应仍调 load_schema（标准路径）"
        assert gate < schema, "分流必须早于 load_schema（否则自定义编号仍 500）"

    def test_router_gate_condition_shape_not_short_circuited(self):
        """🔴 判据必须是**条件形态**，不能只看调用点是否存在。

        `if False and await resolve_is_custom(...)` 会让分流失效，而调用点文本仍在
        ⇒ 只比位置的断言恒绿（变异 M2 实测 GREEN，本条为修正）。
        """
        src = _strip_py_comments(_read(_ROUTER_REL))
        body = _fn_body(src, "export_xlsx")
        assert re.search(
            r"if\s+await\s+resolve_is_custom\(\s*db\s*,\s*working_paper\s*,\s*wp_code\s*\)\s*:",
            body,
        ), "分流条件必须是 `if await resolve_is_custom(db, working_paper, wp_code):`"
        # 不得被 False/None 等常量短路
        assert not re.search(r"if\s+(?:False|None|0)\s+and\s+await\s+resolve_is_custom", body), (
            "分流条件不得被常量短路"
        )

    def test_router_returns_before_schema_for_custom(self):
        """custom 分支必须 return，不能 fall through 到 load_schema。"""
        src = _strip_py_comments(_read(_ROUTER_REL))
        body = _fn_body(src, "export_xlsx")
        gate = _call_pos(body, "resolve_is_custom")
        schema = _call_pos(body, "load_schema")
        seg = body[gate:schema]
        assert "return StreamingResponse" in seg, "custom 分支应直接返回响应"

    def test_resolve_is_custom_three_args(self):
        for rel in (_ROUTER_REL, _ENGINE_REL):
            src = _strip_py_comments(_read(rel))
            for m in re.finditer(r"resolve_is_custom\(([^)]*)\)", src):
                args = [a.strip() for a in m.group(1).split(",") if a.strip()]
                assert len(args) == 3, f"{rel}: resolve_is_custom 应传 3 参，实为 {args}"


# ════════════════════════════════════════════════════════════════════════════
class TestProperty18BothExportPathsBranch:
    """Property 18：两条导出路径都必须分流。"""

    def test_engine_gate_precedes_blank_workbook_fallback(self):
        """🔴 分流必须早于空白 workbook 回退。

        回退分支会吞掉异常并产出「只有 wp_code 一个 sheet 名的空表」却**返回 200**
        —— 比 500 更坏（用户拿到空文件以为导出成功，归档时才发现）。

        🔴 判据已于 2026-08-09 诚实改写：原断言要求源码里存在
        `except (TemplateNotFoundError, Exception)`，而那个元组形态本身就是缺陷
        （元组里带 `Exception` 等于吞一切，且 `TemplateNotFoundError` 是它的子类
        故整个元组是冗余写法）。`wp-export-file-path-resolution` 那轮已把它收敛为
        `except Exception:` 并让回退 workbook **自证失败**。守卫不能继续要求缺陷
        形态存在，否则修好反而打红。现改为：认新形态 + 反向锁死旧形态不得复活。
        """
        src = _strip_py_comments(_read(_ENGINE_REL))
        body = _fn_body(src, "_export_xlsx")
        gate = body.find("if is_custom:")
        fallback = _fallback_except_pos(body)
        assert gate >= 0, "_export_xlsx 应有 custom 分流"
        assert fallback >= 0, "_export_xlsx 的回退分支应仍在（标准路径依赖它）"
        assert gate < fallback, "分流必须早于空白 workbook 回退"

    def test_swallow_all_except_tuple_not_revived(self):
        """🔴 反向锁死：吞一切的元组形态不得复活。

        `except (TemplateNotFoundError, Exception)` 里 `TemplateNotFoundError` 是
        `Exception` 子类 ⇒ 元组等价于裸 `except Exception`，但**读起来像只捕两类**，
        是本轮修掉的可读性陷阱。写成裸 `except Exception:` 才让「这里吞一切」自证。
        """
        src = _strip_py_comments(_read(_ENGINE_REL))
        assert "except (TemplateNotFoundError, Exception)" not in src, (
            "吞一切的元组形态已被收敛为 `except Exception:`，不得复活"
        )

    def test_fallback_workbook_self_evidences_failure(self):
        """🔴 回退 workbook 必须自证是失败，而非「内容本来就空」。

        这是「导出来都是空的」这类用户报告最难排查的一环：HTTP 200 + 一个空 xlsx，
        既看不出失败也看不出原因。故兜底 workbook 第 1 行必须写明「导出失败」+ 原因。
        """
        src = _strip_py_comments(_read(_ENGINE_REL))
        helper = _fn_body(src, "_build_failure_fallback_workbook")
        assert helper, "应有 _build_failure_fallback_workbook 兜底构造器"
        assert "导出失败" in helper, "兜底 workbook 必须写明「导出失败」"
        # 原因文本必须来自真实异常，不能是写死的空话
        assert "type(error).__name__" in helper, "失败原因必须包含异常类型名"
        # 回退分支必须真的调它（防「helper 写了但没接线」）
        body = _fn_body(src, "_export_xlsx")
        assert _call_pos(body, "_build_failure_fallback_workbook") > 0, (
            "_export_xlsx 的回退分支必须调用兜底构造器"
        )

    def test_engine_gate_condition_shape(self):
        """🔴 engine 侧分流同样断言条件形态（防 `if False:` / 常量短路）。"""
        src = _strip_py_comments(_read(_ENGINE_REL))
        body = _fn_body(src, "_export_xlsx")
        assert re.search(r"if\s+is_custom\s*:", body), "分流条件必须是 `if is_custom:`"
        assert not re.search(r"if\s+(?:False|None|0)\s*:", body), (
            "分流条件不得被常量短路"
        )

    def test_engine_custom_branch_calls_dedicated_service(self):
        src = _strip_py_comments(_read(_ENGINE_REL))
        body = _fn_body(src, "_export_xlsx")
        assert _call_pos(body, "export_custom_workpaper") > 0, (
            "_export_xlsx 的 custom 分支应调独立导出服务"
        )

    def test_engine_custom_branch_not_wrapped_in_try(self):
        """CustomExportError 必须冒泡，不得落进空白 workbook 回退。

        🔴 这条曾是**假绿**（2026-08-09 发现）：`fallback` 用旧元组形态 find 得 -1，
        `body[gate:-1]` 仍切出一段几乎完整的正文 ⇒ 断言恰好通过。切片边界为负时
        必须显式判否，不能让 Python 的负索引语义把「判据失效」伪装成「通过」。
        """
        src = _strip_py_comments(_read(_ENGINE_REL))
        body = _fn_body(src, "_export_xlsx")
        gate = body.find("if is_custom:")
        fallback = _fallback_except_pos(body)
        assert gate >= 0 and fallback > gate, (
            f"切片边界无效（gate={gate} fallback={fallback}）—— 判据已失效，勿当通过"
        )
        seg = body[gate:fallback]
        assert "try:" not in seg.split("return export_custom_workpaper")[0], (
            "custom 分流不得被 try 包住（异常必须冒泡）"
        )

    def test_export_single_resolves_custom_and_passes_down(self):
        src = _strip_py_comments(_read(_ENGINE_REL))
        body = _fn_body(src, "export_single")
        assert _call_pos(body, "resolve_is_custom") > 0, "export_single 应判定 custom"
        assert "is_custom=" in body, "判定结果必须传给 _export_xlsx"

    def test_reverse_selfcheck_two_independent_assertions(self):
        """🔴 反向自检：两条路径**各有**断言 —— 只分流一条时另一条必红。

        构造「只分流 /export-xlsx，engine 侧未分流」的替身，
        本类的 engine 断言应判否（证明不是一条断言覆盖两条路径）。
        """
        fake_engine = (
            "    async def _export_xlsx(self, wp, wp_code, schema, meta):\n"
            "        try:\n"
            "            return await export_workpaper_xlsx()\n"
            "        except (TemplateNotFoundError, Exception) as e:\n"
            "            return blank()\n"
        )
        body = _fn_body(fake_engine, "_export_xlsx")
        assert body.find("if is_custom:") < 0, "替身里本就没有分流"
        assert _call_pos(body, "export_custom_workpaper") < 0


# ════════════════════════════════════════════════════════════════════════════
class TestProperty19ExplicitFormat:
    """Property 19：格式判定显式登记 custom。"""

    def test_custom_in_xlsx_types(self):
        from app.services.wp_export.export_engine import _XLSX_TYPES

        assert "custom" in _XLSX_TYPES, "custom 应显式登记为 xlsx 类型"

    def test_why_membership_not_return_value(self):
        """🔴 判据只能靠集合成员，不能靠 determine_export_format 返回值。

        `determine_export_format` 末尾有「默认 xlsx」隐式兜底 ⇒ 即便把 "custom"
        从 `_XLSX_TYPES` 移除，返回值**仍是** "xlsx" ⇒ 按返回值断言的守卫
        对该变异恒绿。这正是「隐式兜底不可断言」的证据。
        """
        from app.services.wp_export import export_engine as ee

        assert ee.determine_export_format(None, "custom") == "xlsx"
        # 未登记的随机类型同样返回 xlsx —— 证明返回值无区分度
        assert ee.determine_export_format(None, "__totally_unknown__") == "xlsx"

    def test_pbt_copy_in_sync(self):
        """🔴 测试侧自带 XLSX_TYPES 副本，必须与生产集合同步。"""
        from app.services.wp_export.export_engine import _XLSX_TYPES

        src = _read("backend/tests/test_pbt_export_format.py")
        m = re.search(r"XLSX_TYPES\s*=\s*\[([^\]]*)\]", src)
        assert m, "未找到测试侧 XLSX_TYPES 副本"
        copy = {x.strip().strip("\"'") for x in m.group(1).split(",") if x.strip()}
        assert copy == set(_XLSX_TYPES), (
            f"测试副本与生产集合不一致: 副本={sorted(copy)} 生产={sorted(_XLSX_TYPES)}"
        )


# ════════════════════════════════════════════════════════════════════════════
class TestProperty20RoundTrip:
    """Property 20：导出字节流可打开，且 cells 逐格相等。"""

    def test_exported_bytes_open_and_cells_match(self, tmp_path):
        import openpyxl

        from app.services.custom_workpaper_projection import project_custom_workpaper

        sheet = "GT-CUSTOM-1"
        cells = {"A1": "项目", "B1": "金额", "A2": "货币资金", "B2": 1234.5}
        fp = _make_xlsx(tmp_path, sheet, cells)

        before = project_custom_workpaper(str(fp), sheet)
        buf = export_custom_workpaper(_FakeWp(file_path=str(fp)))
        assert isinstance(buf, BytesIO)
        data = buf.getvalue()
        assert data, "导出字节流不得为空"

        # ① 可被 openpyxl 打开
        wb = openpyxl.load_workbook(BytesIO(data), data_only=False)
        try:
            assert sheet in wb.sheetnames
        finally:
            wb.close()

        # ② 落盘后重投影，cells 逐格相等（往返可逆）
        out = tmp_path / "roundtrip.xlsx"
        out.write_bytes(data)
        after = project_custom_workpaper(str(out), sheet)
        assert set(after["cells"]) >= set(before["cells"]), "往返后不得丢格"
        for ref, cell in before["cells"].items():
            assert after["cells"][ref]["v"] == cell["v"], f"{ref} 值不一致"

    def test_missing_file_raises_not_blank_workbook(self, tmp_path):
        """🔴 禁 fail-open：文件缺失必须抛错，不得回退空白 workbook。"""
        with pytest.raises(CustomExportError):
            export_custom_workpaper(_FakeWp(file_path=str(tmp_path / "nope.xlsx")))

    def test_empty_path_raises(self):
        with pytest.raises(CustomExportError):
            export_custom_workpaper(_FakeWp(file_path=None))

    def test_corrupt_file_raises(self, tmp_path):
        bad = tmp_path / "bad.xlsx"
        bad.write_bytes(b"not a zip at all")
        with pytest.raises(CustomExportError):
            export_custom_workpaper(_FakeWp(file_path=str(bad)))

    def test_relative_path_resolved_against_backend_root(self, tmp_path, monkeypatch):
        """`file_path` 存的是相对 backend/ 的相对路径（平台既有铁律）。"""
        sheet = "GT-CUSTOM-1"
        fp = _make_xlsx(tmp_path, sheet, {"A1": "x"})
        monkeypatch.chdir(tmp_path)
        # 传相对文件名，CWD 已切到 tmp_path ⇒ 应能解析
        buf = export_custom_workpaper(_FakeWp(file_path=fp.name))
        assert buf.getvalue()

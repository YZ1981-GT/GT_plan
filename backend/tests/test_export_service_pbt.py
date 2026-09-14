"""ExportService 属性测试（advanced-query-module Task 13.4 / 13.5 / 13.6）.

本文件以 Hypothesis 实现 ExportService 的三条正确性属性：

- **Property 12**：导出结构与显示值 round-trip（Validates: Requirements 7.2）。
  随机结果网格（含分组/透视形态、Unicode 标签），导出 ``.xlsx`` 再用 openpyxl 读回，
  断言列标题 / 列顺序 / 行列结构 / 显示值与原结果一致。
- **Property 13**：RFC 5987 文件名编码（Validates: Requirements 7.3）。
  含中文/非 ASCII 显示名，断言 ``Content-Disposition`` 含 ``filename*=UTF-8''`` 且经
  ``urllib.parse.unquote`` 可解码还原为原始显示名。
- **Property 14**：流式与非流式导出一致（Validates: Requirements 8.5，model-based）。
  随机结果集，断言 ``stream_xlsx`` 流式导出的行内容序列与 ``export_xlsx`` 非流式一致。

``export_xlsx`` / ``stream_xlsx`` 为 async，属性体内用 ``asyncio.run`` 驱动。

注：按平台测试提速铁律，每条属性显式使用 ``@settings(max_examples=5)``（优先于 conftest
的 fast profile），保持属性测试稳定且快速。
"""

from __future__ import annotations

import asyncio
import io
from urllib.parse import unquote

from hypothesis import given, settings
from hypothesis import strategies as st
from openpyxl import load_workbook

from app.services.custom_query.export_service import (
    _EMPTY_NOTICE,
    _source_identity_value,
    export_service,
)
from app.services.custom_query.query_orchestrator import ColumnMeta, QueryResult


# ─────────────────────────────────────────────────────────────────────────────
# 生成器：随机结果网格（含 Unicode 标签、可选 addr_id、任意行列形态）
# ─────────────────────────────────────────────────────────────────────────────
# 安全字符集：混合 ASCII / CJK / 希腊/货币符号覆盖 Unicode；剔除控制字符、``=``
# （openpyxl 会把以 ``=`` 开头的串识别为公式）、空白（空串读回为 None 会破坏 round-trip）。
_LABEL_ALPHABET = "abcXYZ0123测试项目金额日本語αΩ€–"

_label_text = st.text(alphabet=_LABEL_ALPHABET, min_size=1, max_size=8)
# 单元格显示值：Unicode 文本 / 整数（限幅避免 xlsx 数值精度问题）/ None（空单元格）
_cell_value = st.one_of(
    st.text(alphabet=_LABEL_ALPHABET, min_size=1, max_size=10),
    st.integers(min_value=-10_000, max_value=10_000),
    st.none(),
)


@st.composite
def _result_grid(draw, *, allow_empty: bool = True):
    """随机 :class:`QueryResult`（列元数据 + 行网格）。

    - 列数 1–4，每列唯一 key ``col{i}``、随机 Unicode 标题；部分列携带唯一 addr_id
      （模拟单源可下钻列，触发「数据来源标识列」导出分支 R7.4）。
    - 行数 0–5（``allow_empty=False`` 时 ≥1），值为 Unicode 文本 / 整数 / 空。
    - 通过变化的行列形态覆盖分组/透视后的任意矩形网格。
    """
    n_cols = draw(st.integers(min_value=1, max_value=4))
    columns: list[ColumnMeta] = []
    for i in range(n_cols):
        addr_id = f"D2/S1/A{i}" if draw(st.booleans()) else None
        columns.append(
            ColumnMeta(
                key=f"col{i}",
                title=draw(_label_text),
                addr_id=addr_id,
                dtype="text",
                semantic_label=draw(st.one_of(st.none(), _label_text)),
            )
        )

    min_rows = 0 if allow_empty else 1
    n_rows = draw(st.integers(min_value=min_rows, max_value=5))
    rows: list[dict] = []
    for _ in range(n_rows):
        rows.append({c.key: draw(_cell_value) for c in columns})

    return QueryResult(columns=columns, rows=rows, total=n_rows)


def _pad_grid(grid: list[tuple], width: int) -> list[tuple]:
    """把每行补齐到 ``width`` 列（read_only 模式会裁掉行尾空单元格，需归一后比对）。"""
    return [row + (None,) * (width - len(row)) for row in grid]


def _read_grid(data: bytes, width: int) -> list[tuple]:
    """用 openpyxl 读回 xlsx 全部行（含表头），补齐行尾空单元格后返回。"""
    wb = load_workbook(io.BytesIO(data), read_only=True)
    ws = wb.active
    grid = [tuple(row) for row in ws.iter_rows(values_only=True)]
    wb.close()
    return _pad_grid(grid, width)


def _grid_width(result: QueryResult) -> int:
    columns = list(result.columns)
    return len(columns) + len([c for c in columns if c.addr_id])


def _expected_grid(result: QueryResult) -> list[tuple]:
    """按 ExportService 契约计算导出后应读回的行网格（含表头 + 数据来源标识列）。"""
    columns = list(result.columns)
    source_cols = [c for c in columns if c.addr_id]
    width = len(columns) + len(source_cols)

    header = [c.title for c in columns] + [
        f"{c.title}·来源(addr_id)" for c in source_cols
    ]
    grid: list[tuple] = [tuple(header)]

    if not result.rows:
        # 0 行 → 仅标题行 + 空提示行（R7.5）
        grid.append((_EMPTY_NOTICE,))
        return _pad_grid(grid, width)

    for row in result.rows:
        values = [row.get(c.key) for c in columns]
        values += [_source_identity_value(c) for c in source_cols]
        grid.append(tuple(values))
    return _pad_grid(grid, width)


# ─────────────────────────────────────────────────────────────────────────────
# Feature: advanced-query-module, Property 12: 导出结构与显示值 round-trip
# ─────────────────────────────────────────────────────────────────────────────
@given(result=_result_grid())
@settings(max_examples=5)
def test_property_12_export_roundtrip_structure_and_values(result: QueryResult) -> None:
    """导出 .xlsx 再读回，列标题/顺序/行列结构/显示值与原结果一致。

    Feature: advanced-query-module, Property 12: 导出结构与显示值 round-trip
    Validates: Requirements 7.2
    """
    data = asyncio.run(export_service.export_xlsx(result))
    read_back = _read_grid(data, _grid_width(result))
    expected = _expected_grid(result)
    assert read_back == expected


# ─────────────────────────────────────────────────────────────────────────────
# Feature: advanced-query-module, Property 13: RFC 5987 文件名编码
# ─────────────────────────────────────────────────────────────────────────────
# 含中文/非 ASCII 的显示名字符集（含空格，验证百分号编码 + 解码 round-trip）。
_FILENAME_ALPHABET = "查询结果报表 abcXYZ_-测试项目日本語αΩ€2025年"


@given(display_name=st.text(alphabet=_FILENAME_ALPHABET, min_size=1, max_size=40))
@settings(max_examples=5)
def test_property_13_rfc5987_filename_encoding(display_name: str) -> None:
    """Content-Disposition 含 filename*=UTF-8'' 且可解码还原为原始显示名。

    Feature: advanced-query-module, Property 13: RFC 5987 文件名编码
    Validates: Requirements 7.3
    """
    header = export_service.build_content_disposition(display_name)

    marker = "filename*=UTF-8''"
    assert marker in header, header

    encoded = header.split(marker, 1)[1]
    decoded = unquote(encoded)

    # 与 ExportService 相同的入口归一：strip 后为空回退为 query.xlsx
    expected = (display_name or "").strip() or "query.xlsx"
    assert decoded == expected


# ─────────────────────────────────────────────────────────────────────────────
# Feature: advanced-query-module, Property 14: 流式与非流式导出一致
# ─────────────────────────────────────────────────────────────────────────────
@given(result=_result_grid())
@settings(max_examples=5)
def test_property_14_stream_matches_materialized(result: QueryResult) -> None:
    """流式导出行内容序列与非流式导出一致（model-based）。

    Feature: advanced-query-module, Property 14: 流式与非流式导出一致
    Validates: Requirements 8.5
    """

    async def _collect_stream() -> bytes:
        chunks: list[bytes] = []
        async for chunk in export_service.stream_xlsx(
            list(result.rows), list(result.columns)
        ):
            chunks.append(chunk)
        return b"".join(chunks)

    materialized = asyncio.run(export_service.export_xlsx(result))
    streamed = asyncio.run(_collect_stream())

    width = _grid_width(result)
    assert _read_grid(streamed, width) == _read_grid(materialized, width)

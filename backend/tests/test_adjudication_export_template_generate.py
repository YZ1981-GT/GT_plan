"""Tests for AdjudicationExportTemplateService.generate() and _build_data_sheet."""

import io
from unittest.mock import AsyncMock

import pytest
from openpyxl import Workbook, load_workbook

from app.services.adjudication_export_template_service import (
    AdjudicationExportTemplateService as Svc,
)


# ── _build_data_sheet tests ──────────────────────────────────────────────────


class TestBuildDataSheet:
    """Sub-tasks 5.1 + 5.2: title row, multi-row header, and row skeleton."""

    def _make_ws(self):
        wb = Workbook()
        return wb.active

    def test_title_row_content(self):
        ws = self._make_ws()
        Svc._build_data_sheet(ws, "D2-1", "应收账款", [])
        assert ws.cell(1, 1).value == "应收账款审定表 D2-1"

    def test_title_row_bold(self):
        ws = self._make_ws()
        Svc._build_data_sheet(ws, "D2-1", "应收账款", [])
        assert ws.cell(1, 1).font.bold is True

    def test_title_row_font_size(self):
        ws = self._make_ws()
        Svc._build_data_sheet(ws, "D2-1", "应收账款", [])
        assert ws.cell(1, 1).font.size == 14

    def test_title_merge_no_aging(self):
        ws = self._make_ws()
        Svc._build_data_sheet(ws, "D1-1", "货币资金", [])
        # 12 columns for non-aging
        merged = [str(m) for m in ws.merged_cells.ranges]
        assert "A1:L1" in merged  # A=1, L=12

    def test_title_merge_with_aging(self):
        ws = self._make_ws()
        aging = ["1年以内", "1-2年", "2-3年"]
        Svc._build_data_sheet(ws, "D2-1", "应收账款", [], aging_headers=aging)
        # 12 + 3 = 15 columns → A1:O1
        merged = [str(m) for m in ws.merged_cells.ranges]
        assert "A1:O1" in merged  # O = column 15

    def test_header_rows_present(self):
        ws = self._make_ws()
        Svc._build_data_sheet(ws, "D1-1", "货币资金", [])
        assert ws.cell(2, 1).value == "项目"
        assert ws.cell(2, 2).value == "期初"
        assert ws.cell(2, 6).value == "期末"
        assert ws.cell(3, 2).value == "未审"

    def test_row_skeleton_basic(self):
        ws = self._make_ws()
        skeleton = [
            {"id": "1", "item": "应收账款", "is_section": False, "is_total": False, "indent": 0},
            {"id": "2", "item": "其他应收款", "is_section": False, "is_total": False, "indent": 0},
        ]
        Svc._build_data_sheet(ws, "D2-1", "应收账款", skeleton)
        assert ws.cell(4, 1).value == "应收账款"
        assert ws.cell(5, 1).value == "其他应收款"

    def test_row_skeleton_section_bold(self):
        ws = self._make_ws()
        skeleton = [
            {"id": "1", "item": "流动资产", "is_section": True, "is_total": False, "indent": 0},
        ]
        Svc._build_data_sheet(ws, "D2-1", "应收账款", skeleton)
        assert ws.cell(4, 1).font.bold is True

    def test_row_skeleton_total_bold(self):
        ws = self._make_ws()
        skeleton = [
            {"id": "1", "item": "合计", "is_section": False, "is_total": True, "indent": 0},
        ]
        Svc._build_data_sheet(ws, "D2-1", "应收账款", skeleton)
        assert ws.cell(4, 1).font.bold is True

    def test_row_skeleton_normal_not_bold(self):
        ws = self._make_ws()
        skeleton = [
            {"id": "1", "item": "明细行", "is_section": False, "is_total": False, "indent": 0},
        ]
        Svc._build_data_sheet(ws, "D2-1", "应收账款", skeleton)
        # Normal rows should not have bold font explicitly set
        cell = ws.cell(4, 1)
        assert cell.font.bold is not True

    def test_empty_skeleton_no_data_rows(self):
        ws = self._make_ws()
        Svc._build_data_sheet(ws, "D1-1", "货币资金", [])
        # Row 4 should be empty
        assert ws.cell(4, 1).value is None


# ── generate() tests ─────────────────────────────────────────────────────────


class TestGenerate:
    """Sub-task 5.3: generate() orchestrator returns (BytesIO, filename)."""

    @pytest.fixture
    def mock_db(self):
        return AsyncMock()

    @pytest.mark.asyncio
    async def test_returns_bytesio_and_filename(self, mock_db):
        buf, filename = await Svc.generate(
            wp_id="wp-123",
            wp_code="D1-1",
            wp_name="货币资金",
            db=mock_db,
            template_file_path=None,
        )
        assert isinstance(buf, io.BytesIO)
        assert filename == "D1-1_货币资金审定表_模板.xlsx"

    @pytest.mark.asyncio
    async def test_workbook_has_two_sheets(self, mock_db):
        buf, _ = await Svc.generate(
            wp_id="wp-123",
            wp_code="D1-1",
            wp_name="货币资金",
            db=mock_db,
            template_file_path=None,
        )
        wb = load_workbook(buf)
        assert len(wb.sheetnames) == 2
        assert wb.sheetnames[0] == "编制说明"
        assert wb.sheetnames[1] == "货币资金审定表"

    @pytest.mark.asyncio
    async def test_data_sheet_title(self, mock_db):
        buf, _ = await Svc.generate(
            wp_id="wp-123",
            wp_code="F3-1",
            wp_name="固定资产",
            db=mock_db,
            template_file_path=None,
        )
        wb = load_workbook(buf)
        ws = wb["固定资产审定表"]
        assert ws.cell(1, 1).value == "固定资产审定表 F3-1"

    @pytest.mark.asyncio
    async def test_non_aging_no_db_call(self, mock_db):
        """D1-1 is not in AGING_SUBJECTS, so no DB resolution should occur."""
        await Svc.generate(
            wp_id="wp-123",
            wp_code="D1-1",
            wp_name="货币资金",
            db=mock_db,
            template_file_path=None,
        )
        # _resolve_aging_headers should not be called
        mock_db.assert_not_awaited()

    @pytest.mark.asyncio
    async def test_bytesio_is_valid_xlsx(self, mock_db):
        buf, _ = await Svc.generate(
            wp_id="wp-123",
            wp_code="D1-1",
            wp_name="货币资金",
            db=mock_db,
            template_file_path=None,
        )
        # Should not raise
        wb = load_workbook(buf)
        assert wb is not None

    @pytest.mark.asyncio
    async def test_filename_format_chinese(self, mock_db):
        _, filename = await Svc.generate(
            wp_id="wp-123",
            wp_code="K1-1",
            wp_name="短期借款",
            db=mock_db,
            template_file_path=None,
        )
        assert filename == "K1-1_短期借款审定表_模板.xlsx"

    @pytest.mark.asyncio
    async def test_empty_template_path_produces_empty_skeleton(self, mock_db):
        buf, _ = await Svc.generate(
            wp_id="wp-123",
            wp_code="D1-1",
            wp_name="货币资金",
            db=mock_db,
            template_file_path=None,
        )
        wb = load_workbook(buf)
        ws = wb["货币资金审定表"]
        # No skeleton → row 4 col 1 should be empty
        assert ws.cell(4, 1).value is None

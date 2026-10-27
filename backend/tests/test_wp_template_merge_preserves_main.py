"""多文件底稿合并不再整册重存主文件（spec post-push-residual-defect-closure R4.5）。

按编码生成 / 在线编辑首开都经 ``init_workpaper_from_template``：多文件底稿（C2~C15 带「评价控制偏差」、E1 / D2 / D4 / F2 …）
把其余文件的页按值追加到主文件。原先 openpyxl 打开主文件、建页写值后整册重存 ⇒ 主文件的 x14 数据验证、公式缓存值、
条件格式扩展等全丢（C2 的 3 条数据验证生成后为 0）。现在只追加部件，主文件原有部件除工作簿目录三件外逐字节不变；
追加出来的页与原先 openpyxl 写法逐页逐格相同。
"""
from __future__ import annotations

import uuid
import zipfile
from pathlib import Path

import openpyxl
import pytest

from app.services import wp_template_init_service as init
from app.services import xlsx_cell_patch as xp

#: 追加页会改的部件（styles.xml 只在追加页里有日期格时改）
_ALLOWED = {"xl/workbook.xml", "xl/_rels/workbook.xml.rels", "[Content_Types].xml", "xl/styles.xml"}


def _parts(path: Path) -> dict[str, bytes]:
    with zipfile.ZipFile(path) as zf:
        return {n: zf.read(n) for n in zf.namelist()}


def _dv(parts: dict[str, bytes]) -> int:
    return sum(v.count(b"<x14:dataValidation ") for n, v in parts.items() if n.startswith("xl/worksheets/"))


def _generate(tmp_path: Path, monkeypatch, code: str, sub: str) -> Path:
    monkeypatch.setattr(init, "STORAGE_DIR", tmp_path / sub)
    pid = uuid.uuid4()
    (tmp_path / sub / str(pid)).mkdir(parents=True)
    path = init.init_workpaper_from_template(pid, uuid.uuid4(), code)
    assert path is not None, f"{code} 找不到模板"
    return Path(path)


def _values(path: Path, titles: list[str]) -> dict[str, dict[str, object]]:
    wb = openpyxl.load_workbook(path)
    try:
        return {t: {c.coordinate: c.value for row in wb[t].iter_rows() for c in row if c.value is not None}
                for t in titles}
    finally:
        wb.close()


@pytest.mark.parametrize("code", ["C2", "C14", "E1", "D2", "F2"])
def test_multi_file_merge_keeps_main_template_and_matches_old_values(tmp_path, monkeypatch, code):
    main = init.find_template_file_any(code)
    others = [f for f in init.find_all_template_files(code) if f != main]
    assert others, f"{code} 应是多文件底稿（模板目录变了？）"
    template = _parts(main)

    path = _generate(tmp_path, monkeypatch, code, "new")
    out = _parts(path)
    changed = sorted(n for n in template if template[n] != out.get(n))
    assert set(changed) <= _ALLOWED, f"{code} 主文件原有部件被改：{[n for n in changed if n not in _ALLOWED]}"
    assert _dv(out) == _dv(template), f"{code} 数据验证条数变了"
    names = xp.workbook_sheet_names(path)
    main_names = xp.workbook_sheet_names(main)
    assert names[: len(main_names)] == main_names
    merged = names[len(main_names):]
    assert merged, f"{code} 没有追加任何页"

    # 追加页的值 == 源文件同名页的值（公式取缓存结果；独立读源文件，不经被测的合并代码）
    source: dict[str, dict[str, object]] = {}
    for f in others:
        wb = openpyxl.load_workbook(f, read_only=True, data_only=True)
        try:
            for title in wb.sheetnames:
                if title[:31] in merged and title[:31] not in source:
                    source[title[:31]] = {c.coordinate: c.value for row in wb[title].iter_rows() for c in row
                                          if getattr(c, "value", None) is not None}
        finally:
            wb.close()
    same_named = [t for t in merged if t in source]
    assert same_named, f"{code} 追加页在源文件里一页都找不到"
    got = _values(path, same_named)
    for title in same_named:
        assert got[title] == source[title], f"{code}「{title}」追加页的值与源文件不同"

    # 与原先 openpyxl 写法（退路）逐页逐格对比
    def _raise(*_a, **_k):
        raise xp.XlsxPatchError("测试：走退路")

    monkeypatch.setattr(xp, "append_value_sheets", _raise)
    old = _generate(tmp_path, monkeypatch, code, "old")
    old_names = xp.workbook_sheet_names(old)
    assert old_names == names, f"{code} 页签与原先写法不同"
    assert _values(path, merged) == _values(old, merged), f"{code} 追加页的值与原先写法不同"


def test_single_file_template_is_copied_byte_identical(tmp_path, monkeypatch):
    path = _generate(tmp_path, monkeypatch, "H4", "x")
    assert path.read_bytes() == init.find_template_file_any("H4").read_bytes()

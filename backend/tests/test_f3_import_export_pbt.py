"""F3 导入导出 Round-Trip PBT — Task 10.1

生成随机行 → export xlsx → import → 验证等价。
"""

from __future__ import annotations

import io

from hypothesis import given, settings as hyp_settings
from hypothesis import strategies as st
from openpyxl import Workbook, load_workbook

from app.routers.wp_render_strategies._cycle_import_export_common import (
    export_row_by_keys,
    parse_row_by_headers,
)
from app.routers.wp_render_strategies._f3_import_export import _F3_SPECS, _normalize_rows

st_float = st.floats(min_value=0, max_value=1e8, allow_nan=False, allow_infinity=False)
st_text = st.text(min_size=0, max_size=20, alphabet=st.characters(categories=("L", "N")))

st_f3_7_credit_row = st.fixed_dictionaries({
    "rowId": st.uuids().map(str),
    "seq": st.integers(min_value=1, max_value=999),
    "summary": st_text,
    "counterAccount": st_text,
    "amount": st_float,
    "voucherDate": st.just("2025-03-01"),
    "voucherNo": st.text(min_size=1, max_size=12, alphabet=st.characters(categories=("L", "N"))),
    "noteType": st.sampled_from(["银行承兑", "商业承兑", ""]),
    "acceptor": st_text,
    "purchaseContractCheck": st.sampled_from(["", "Y", "N"]),
    "goodsReceiptCheck": st.sampled_from(["", "Y", "N"]),
    "auditConclusion": st_text,
    "remark": st_text,
})

st_f3_4_row = st.fixed_dictionaries({
    "rowId": st.uuids().map(str),
    "seq": st.integers(min_value=1, max_value=99),
    "drawer": st_text,
    "faceValue": st_float,
    "interestRate": st.floats(min_value=0, max_value=20, allow_nan=False, allow_infinity=False),
    "issueDate": st.just("2025-01-01"),
    "dueDate": st.just("2025-07-01"),
    "termDays": st.integers(min_value=0, max_value=365),
    "interestStart": st.just("2025-01-01"),
    "interestEnd": st.just("2025-12-31"),
    "accruedDays": st.integers(min_value=0, max_value=365),
    "payableInterest": st_float,
    "bookInterest": st_float,
    "variance": st_float,
})


def _roundtrip(sheet: str, rows: list[dict]) -> list[dict]:
    sp = _F3_SPECS[sheet]
    headers = sp["headers"]
    keys = sp["field_keys"]

    wb = Workbook()
    ws = wb.active
    ws.title = sheet
    ws.append(headers)
    for row in rows:
        ws.append(export_row_by_keys(row, keys))

    buf = io.BytesIO()
    wb.save(buf)
    buf.seek(0)

    wb2 = load_workbook(buf, read_only=True, data_only=True)
    ws2 = wb2.active
    actual_headers = [
        str(c.value).strip() if c.value else ""
        for c in next(ws2.iter_rows(min_row=1, max_row=1))
    ]
    imported: list[dict] = []
    for row_tuple in ws2.iter_rows(min_row=2, values_only=True):
        if all(v is None for v in row_tuple):
            continue
        imported.append(parse_row_by_headers(row_tuple, actual_headers, keys))
    wb2.close()
    return _normalize_rows(imported)


def _assert_close(actual: float, expected: float, msg: str) -> None:
    diff = abs(float(actual or 0) - float(expected or 0))
    assert diff < 1e-4 * max(abs(float(actual or 0)), abs(float(expected or 0)), 1.0) + 1e-6, msg


@hyp_settings(max_examples=5)
@given(rows=st.lists(st_f3_7_credit_row, min_size=1, max_size=5))
def test_f3_7_credit_export_import_roundtrip(rows: list[dict]) -> None:
    imported = _roundtrip("F3-7-credit", rows)
    assert len(imported) == len(rows)
    for i, (orig, imp) in enumerate(zip(rows, imported)):
        assert imp["summary"] == orig["summary"], f"row {i} summary"
        assert imp["voucherNo"] == orig["voucherNo"], f"row {i} voucherNo"
        _assert_close(imp["amount"], orig["amount"], f"row {i} amount")
        assert imp.get("rowId"), f"row {i} should have rowId after normalize"


@hyp_settings(max_examples=5)
@given(rows=st.lists(st_f3_4_row, min_size=1, max_size=5))
def test_f3_4_export_import_roundtrip(rows: list[dict]) -> None:
    imported = _roundtrip("F3-4", rows)
    assert len(imported) == len(rows)
    for i, (orig, imp) in enumerate(zip(rows, imported)):
        assert imp["drawer"] == orig["drawer"], f"row {i} drawer"
        _assert_close(imp["faceValue"], orig["faceValue"], f"row {i} faceValue")
        _assert_close(imp["interestRate"], orig["interestRate"], f"row {i} interestRate")


def test_f3_import_normalize_adds_row_id() -> None:
    rows = [{"seq": 1, "summary": "测试", "amount": 100}]
    normalized = _normalize_rows(rows)
    assert normalized[0].get("rowId")
    assert normalized[0]["seq"] == 1

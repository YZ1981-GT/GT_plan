"""E0 隐藏 sheet 必须 skip —— 以源 xlsx `sheet_state` 为唯一裁决者。

背景（2026-08-02）
-----------------
用户反馈 E0 底稿多出 `回函情况汇编` / `银行函证其他信息核对表E0-5` /
`邮件传真回函核对记录F1-12` 三个页签，而 WPS 打开源模板看不到它们。实测：

1. 这三张在源 xlsx 里 ``sheet_state == 'hidden'``；
2. `底稿目录` 的索引清单只列 9 项（E0A + E0-1 ~ E0-8），也不含它们。

两个独立证据一致 → **不属于 E0 底稿集合**，按平台既有机制（`wp_code_overrides.json`
的 ``skip``，E0 另 7 张遗留 sheet 早已如此）从 render-config 剔除。

根因（平台级，本文件不修）
--------------------------
``backend/scripts/analyze/analyze_wp_templates.py`` 枚举 ``wb.sheetnames`` 时不判
``sheet_state``，隐藏 sheet 经 `workpaper_template_analysis.json` →
`seed_workpaper_sheet_classification.py` → `workpaper_sheet_classification` →
render-config → 前端页签。全库 351 模板 / 2722 sheet 中 247 张隐藏（180 个模板），
去歧义后 353 行 / 228 个 wp_code 受影响。本文件只钉死 E0 一册；平台级收口另立 spec。

守卫形态
--------
复刻 ``wp_render_config.py`` 的三段 skip 判定（完整 sheet_name / 尾码 / 首码），
对 E0 源模板逐 sheet 双向断言：

- ``hidden`` → 必须被 skip（漏配即红）
- ``visible`` → 必须**不**被 skip（误伤即红；`银行函证其他信息核对表E0-5` 与真实的
  `应付银行承兑汇票发函记录表E0-5` 同尾码 `E0-5`，若图省事把 `E0-5` 标 skip
  会连真表一起杀掉 —— 这条断言就是拦它的）
"""
from __future__ import annotations

import re
from pathlib import Path

import openpyxl
import pytest

from app.routers.wp_render_config import _SHEET_CODE_RE
from app.services.wp_classification_service import refresh_wp_code_overrides

_REPO_ROOT = Path(__file__).resolve().parents[2]
_E0_XLSX = (
    _REPO_ROOT / "backend" / "wp_templates" / "E"
    / "E0 货币资金 - 函证（Leap应对措施-函证）.xlsx"
)

# 源模板实测：20 张 sheet 中 10 隐藏 10 可见（反向自检的锚，防路径错→空集合→断言恒真）
_EXPECTED_TOTAL = 20
_EXPECTED_HIDDEN = 10


@pytest.fixture(scope="module")
def overrides() -> dict[str, str]:
    return refresh_wp_code_overrides()


@pytest.fixture(scope="module")
def sheet_states() -> list[tuple[str, str]]:
    """[(sheet_name, sheet_state)]，顺序同源 xlsx。"""
    assert _E0_XLSX.exists(), f"源模板不存在：{_E0_XLSX}"
    wb = openpyxl.load_workbook(_E0_XLSX)
    try:
        return [(name, wb[name].sheet_state) for name in wb.sheetnames]
    finally:
        wb.close()


def is_skipped(sheet_name: str, ovr: dict[str, str]) -> bool:
    """复刻 wp_render_config.py L708 / L722 / L727 的三段 skip 判定。"""
    # 1) 完整 sheet_name 精确匹配
    if ovr.get(sheet_name) == "skip":
        return True
    # 2) 尾部编码
    m = _SHEET_CODE_RE.search(sheet_name)
    if m and ovr.get(m.group(1)) == "skip":
        return True
    # 3) 编码在开头
    if not m:
        m2 = re.match(r"([A-Z]\d+(?:-\d+)*)", sheet_name)
        if m2 and ovr.get(m2.group(1)) == "skip":
            return True
    return False


class TestSourceTemplateFacts:
    """反向自检：源模板真的读到了内容，且隐藏/可见分布符合实测。"""

    def test_sheet_counts(self, sheet_states):
        assert len(sheet_states) == _EXPECTED_TOTAL
        hidden = [n for n, st in sheet_states if st != "visible"]
        assert len(hidden) == _EXPECTED_HIDDEN, hidden

    def test_three_reported_sheets_are_hidden(self, sheet_states):
        """用户报的三张确为 hidden（本 spec 的事实基础，写错名字必红）。"""
        states = dict(sheet_states)
        for name in (
            "回函情况汇编",
            "银行函证其他信息核对表E0-5",
            "邮件传真回函核对记录F1-12",
        ):
            assert name in states, f"源 xlsx 无此 tab：{name}"
            assert states[name] == "hidden", f"{name} 实为 {states[name]}"

    def test_directory_index_lists_only_nine_items(self):
        """`底稿目录` 索引只列 9 项（E0A + E0-1~E0-8）—— 第二个独立证据。"""
        wb = openpyxl.load_workbook(_E0_XLSX, data_only=True)
        try:
            ws = wb["底稿目录"]
            codes = [
                str(ws.cell(row=r, column=6).value or "").strip()
                for r in range(3, 12)
            ]
        finally:
            wb.close()
        assert codes == [
            "E0A", "E0-1", "E0-2", "E0-3", "E0-4", "E0-5", "E0-6", "E0-7", "E0-8",
        ], codes


class TestHiddenSheetsSkipped:
    def test_every_hidden_sheet_is_skipped(self, sheet_states, overrides):
        leaked = [
            name for name, st in sheet_states
            if st != "visible" and not is_skipped(name, overrides)
        ]
        assert not leaked, (
            f"E0 源模板的隐藏 sheet 未被 skip，会渲染成多余页签：{leaked}。"
            "请在 backend/app/data/wp_code_overrides.json 按**完整 sheet_name** 标 skip。"
        )

    def test_no_visible_sheet_is_skipped(self, sheet_states, overrides):
        killed = [
            name for name, st in sheet_states
            if st == "visible" and is_skipped(name, overrides)
        ]
        assert not killed, (
            f"E0 真实底稿被 skip 误伤：{killed}。"
            "常见成因 = 按尾部编码（如 E0-5）标 skip 而非完整 sheet_name。"
        )

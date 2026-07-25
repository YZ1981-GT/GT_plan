"""内容控件注入 × 标记清理集成测试 — deliverable-lineage-content-control Task 3.1.

验证 inject_content_controls_for_blocks（包内部内容）与 remove_section_markers 的
协作顺序正确（Property 2 bookmark 并存 / Property 4 内容不变 / 标记仍被清理）。
不依赖 DB / OnlyOffice，用合成 docx 模拟 note 导出 step6→step7 之间的状态。
"""

from __future__ import annotations

from docx import Document
from docx.oxml import OxmlElement
from docx.oxml.ns import qn

from app.services.content_control_injector import inject_content_controls_for_blocks
from app.services.section_anchor_utils import anchor_name
from app.services.word_doc_utils import remove_section_markers, scan_section_blocks


def _build_doc():
    doc = Document()
    doc.add_paragraph("##SECTION:八、1##")
    p_a = doc.add_paragraph("货币资金内容 A")
    doc.add_paragraph("货币资金内容 B")
    doc.add_paragraph("##/SECTION:八、1##")
    doc.add_paragraph("##SECTION:五、1##")
    doc.add_paragraph("应收账款内容")
    doc.add_paragraph("##/SECTION:五、1##")
    # 在内容段落加一个 note_section bookmark（模拟 note 导出的书签，验证并存不被破坏）
    _add_bookmark(p_a, "note_section_八、1")
    return doc


def _add_bookmark(paragraph, name: str):
    start = OxmlElement("w:bookmarkStart")
    start.set(qn("w:id"), "5001")
    start.set(qn("w:name"), name)
    paragraph._p.insert(0, start)
    end = OxmlElement("w:bookmarkEnd")
    end.set(qn("w:id"), "5001")
    paragraph._p.append(end)


def _body_marker_texts(doc) -> list[str]:
    """body 级段落中仍为 SECTION 标记的文本。"""
    out = []
    for child in doc.element.body.iterchildren():
        if child.tag == qn("w:p"):
            txt = "".join(t.text or "" for t in child.iter(qn("w:t")))
            if txt.strip().startswith("##SECTION") or txt.strip().startswith("##/SECTION"):
                out.append(txt.strip())
    return out


def _sdt_tags(doc) -> list[str]:
    tags = []
    for sdt in doc.element.body.iter(qn("w:sdt")):
        pr = sdt.find(qn("w:sdtPr"))
        tag_el = pr.find(qn("w:tag")) if pr is not None else None
        if tag_el is not None:
            tags.append(tag_el.get(qn("w:val")))
    return tags


def _all_text(doc) -> list[str]:
    return [
        "".join(t.text or "" for t in p.iter(qn("w:t")))
        for p in doc.element.body.iter(qn("w:p"))
    ]


def _bookmark_names(doc) -> list[str]:
    return [
        bm.get(qn("w:name"))
        for bm in doc.element.body.iter(qn("w:bookmarkStart"))
    ]


def test_inject_keeps_markers_body_level_then_removed():
    doc = _build_doc()
    before_text = _all_text(doc)

    blocks = scan_section_blocks(doc)
    count = inject_content_controls_for_blocks(doc, blocks)
    assert count == 2

    # 注入后：sdt 带正确 Tag；开闭标记仍在 body 级（供 remove_section_markers 清理）
    tags = _sdt_tags(doc)
    assert anchor_name("八、1") in tags
    assert anchor_name("五、1") in tags
    assert len(_body_marker_texts(doc)) == 4  # 两节各 open+close
    # 可见文本与顺序不变（Property 4）
    assert _all_text(doc) == before_text
    # note_section bookmark 随内容移入 sdtContent，未被破坏（Property 2 并存）
    assert "note_section_八、1" in _bookmark_names(doc)

    # step 7：清理标记 → body 级标记消失，sdt + 内容存活
    removed = remove_section_markers(doc)
    assert removed >= 4
    assert _body_marker_texts(doc) == []
    assert anchor_name("八、1") in _sdt_tags(doc)
    remaining = _all_text(doc)
    assert "货币资金内容 A" in remaining
    assert "应收账款内容" in remaining
    assert "note_section_八、1" in _bookmark_names(doc)


def test_empty_section_skipped():
    doc = Document()
    doc.add_paragraph("##SECTION:六、9##")
    doc.add_paragraph("##/SECTION:六、9##")  # 无内部内容
    blocks = scan_section_blocks(doc)
    assert inject_content_controls_for_blocks(doc, blocks) == 0
    assert _sdt_tags(doc) == []

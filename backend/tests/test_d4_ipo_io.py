"""D4-29/30/31/32 IPO 舞弊应对四表 IO 行为守卫。

验证导出→导入往返后数据结构保真：
- D4-29 客户信息检查表：{id,name,fields} 转置表
- D4-30 客户访谈记录汇总表：{customers, customDimensions}
- D4-31 客户访谈记录：单对象问卷，q1_relation 是 string[]
- D4-32 资金流水检查：六组分组 [{key, rows}]
"""
from __future__ import annotations

import json
import pytest
from uuid import uuid4

from app.routers.wp_render_strategies._d4_import_export import (
    _parse_d4_29_rows,
    _export_d4_29_rows,
    _parse_d4_30_rows,
    _export_d4_30_rows,
    _parse_d4_31_row,
    _export_d4_31_row,
    _parse_d4_32_rows,
    _export_d4_32_rows,
    _SPECIAL_ITEM_IDS,
    _D4_29_HEADER_TO_KEY,
    _D4_30_HEADER_TO_KEY,
    _D4_31_HEADER_TO_KEY,
    _D4_32_HEADER_TO_KEY,
    _D4_32_GROUP_KEYS,
    _D4_32_GROUP_LABELS,
    _SHEET_HEADERS,
    _resolve_item_id,
)

# ═══════════════════════════════════════════════════════════════════════════════
# 1. item_id 映射正确性
# ═══════════════════════════════════════════════════════════════════════════════


class TestItemIdMapping:
    """Req 1.3: item_id 必须与前端一致，否则导入写孤儿、导出导空。"""

    def test_d4_29_item_id(self):
        assert _SPECIAL_ITEM_IDS["D4-29"] == "D4-29-customers"
        assert _resolve_item_id("D4-29") == "D4-29-customers"

    def test_d4_30_item_id(self):
        assert _SPECIAL_ITEM_IDS["D4-30"] == "D4-30-customers"
        assert _resolve_item_id("D4-30") == "D4-30-customers"

    def test_d4_31_item_id(self):
        assert _SPECIAL_ITEM_IDS["D4-31"] == "D4-31-interview"
        assert _resolve_item_id("D4-31") == "D4-31-interview"

    def test_d4_32_item_id(self):
        assert _SPECIAL_ITEM_IDS["D4-32"] == "D4-32-groups"
        assert _resolve_item_id("D4-32") == "D4-32-groups"


# ═══════════════════════════════════════════════════════════════════════════════
# 2. 列头声明完整性
# ═══════════════════════════════════════════════════════════════════════════════


class TestHeaderDeclarations:
    """Req 1.1: 四表列头必须声明且映射完备。"""

    def test_d4_29_headers_exist(self):
        h = _SHEET_HEADERS["D4-29"]
        assert "客户名称" in h
        assert "统一社会信用代码" in h
        assert len(h) >= 20

    def test_d4_30_headers_exist(self):
        h = _SHEET_HEADERS["D4-30"]
        assert "客户名称" in h
        assert "访谈时间" in h
        assert "访谈结论" in h

    def test_d4_31_headers_exist(self):
        h = _SHEET_HEADERS["D4-31"]
        assert "访谈对象" in h
        assert "客户业务关系" in h
        assert "签字日期" in h

    def test_d4_32_headers_exist(self):
        h = _SHEET_HEADERS["D4-32"]
        assert "单位名称/姓名" in h
        assert "是否发现异常交易" in h

    def test_d4_29_headers_mapped(self):
        """每个列头必须在映射表中有对应英文 key。"""
        for h in _SHEET_HEADERS["D4-29"]:
            assert h in _D4_29_HEADER_TO_KEY, f"D4-29 列头 '{h}' 未映射"

    def test_d4_30_headers_mapped(self):
        for h in _SHEET_HEADERS["D4-30"]:
            assert h in _D4_30_HEADER_TO_KEY, f"D4-30 列头 '{h}' 未映射"

    def test_d4_31_headers_mapped(self):
        for h in _SHEET_HEADERS["D4-31"]:
            assert h in _D4_31_HEADER_TO_KEY, f"D4-31 列头 '{h}' 未映射"

    def test_d4_32_headers_mapped(self):
        for h in _SHEET_HEADERS["D4-32"]:
            assert h in _D4_32_HEADER_TO_KEY, f"D4-32 列头 '{h}' 未映射"


# ═══════════════════════════════════════════════════════════════════════════════
# 3. D4-29 客户信息检查表 IO 往返
# ═══════════════════════════════════════════════════════════════════════════════


class TestD429Roundtrip:
    """Req 1.1/1.2: D4-29 转置表 export→import 往返保真。"""

    HEADERS = _SHEET_HEADERS["D4-29"]

    def _make_customer(self, name: str, **fields) -> dict:
        return {"id": f"cust-{uuid4().hex[:12]}", "name": name, "fields": fields}

    def test_roundtrip_single_customer(self):
        cust = self._make_customer("客户甲", creditCode="91110000MA01B", regAddress="北京市朝阳区", isRelated="是")
        from openpyxl import Workbook
        wb = Workbook(); ws = wb.active
        ws.append(self.HEADERS)
        _export_d4_29_rows(ws, [cust], self.HEADERS)
        # 读回
        data_rows = list(ws.iter_rows(min_row=2, values_only=True))
        result = _parse_d4_29_rows(data_rows, self.HEADERS)
        assert len(result) == 1
        r = result[0]
        assert r["name"] == "客户甲"
        assert r["fields"]["creditCode"] == "91110000MA01B"
        assert r["fields"]["regAddress"] == "北京市朝阳区"
        assert r["fields"]["isRelated"] == "是"

    def test_roundtrip_multiple_customers(self):
        custs = [
            self._make_customer("客户A", legalRep="张三"),
            self._make_customer("客户B", isBlacklisted="否", bizStatus="正常"),
        ]
        from openpyxl import Workbook
        wb = Workbook(); ws = wb.active
        ws.append(self.HEADERS)
        _export_d4_29_rows(ws, custs, self.HEADERS)
        data_rows = list(ws.iter_rows(min_row=2, values_only=True))
        result = _parse_d4_29_rows(data_rows, self.HEADERS)
        assert len(result) == 2
        assert result[0]["name"] == "客户A"
        assert result[0]["fields"]["legalRep"] == "张三"
        assert result[1]["name"] == "客户B"
        assert result[1]["fields"]["isBlacklisted"] == "否"

    def test_empty_fields_preserved(self):
        cust = self._make_customer("空客户")
        from openpyxl import Workbook
        wb = Workbook(); ws = wb.active
        ws.append(self.HEADERS)
        _export_d4_29_rows(ws, [cust], self.HEADERS)
        data_rows = list(ws.iter_rows(min_row=2, values_only=True))
        result = _parse_d4_29_rows(data_rows, self.HEADERS)
        assert len(result) == 1
        assert result[0]["name"] == "空客户"
        # 所有字段为空串
        for key in _D4_29_HEADER_TO_KEY.values():
            if key != "_name":
                assert result[0]["fields"][key] == ""

    def test_skip_empty_name(self):
        """无客户名的行应被跳过。"""
        data_rows = [(None,) * len(self.HEADERS)]
        result = _parse_d4_29_rows(data_rows, self.HEADERS)
        assert len(result) == 0


# ═══════════════════════════════════════════════════════════════════════════════
# 4. D4-30 客户访谈汇总表 IO 往返
# ═══════════════════════════════════════════════════════════════════════════════


class TestD430Roundtrip:
    """Req 1.3: D4-30 保留转置矩阵和 customDimensions。"""

    HEADERS = _SHEET_HEADERS["D4-30"]

    def test_roundtrip_with_custom_dimensions(self):
        payload = {
            "customers": [
                {"id": "iv-abc123", "name": "客户X", "fields": {"time": "2025-06-01", "reason": "前十大", "custom_dim1": "自定义值"}},
            ],
            "customDimensions": [{"key": "custom_dim1", "label": "自定义维度1"}],
        }
        from openpyxl import Workbook
        wb = Workbook(); ws = wb.active
        ws.append(self.HEADERS)
        _export_d4_30_rows(ws, payload, self.HEADERS)
        # 读回（含追加的自定义列头）
        actual_headers = [cell.value for cell in next(ws.iter_rows(min_row=1, max_row=1))]
        actual_headers = [str(h).strip() if h else "" for h in actual_headers]
        data_rows = list(ws.iter_rows(min_row=2, values_only=True))
        result = _parse_d4_30_rows(data_rows, actual_headers)
        assert len(result["customers"]) == 1
        c = result["customers"][0]
        assert c["name"] == "客户X"
        assert c["fields"]["time"] == "2025-06-01"
        assert c["fields"]["reason"] == "前十大"
        # 自定义维度恢复
        assert len(result["customDimensions"]) >= 1
        # 自定义维度值保留（key 可能重生成，但 label 和值一致）
        dim_labels = {d["label"] for d in result["customDimensions"]}
        assert "自定义维度1" in dim_labels

    def test_roundtrip_no_custom(self):
        payload = {
            "customers": [
                {"id": "iv-xyz", "name": "客户Y", "fields": {"conclusion": "正常", "indexRef": "D4-31-1"}},
            ],
            "customDimensions": [],
        }
        from openpyxl import Workbook
        wb = Workbook(); ws = wb.active
        ws.append(self.HEADERS)
        _export_d4_30_rows(ws, payload, self.HEADERS)
        data_rows = list(ws.iter_rows(min_row=2, values_only=True))
        result = _parse_d4_30_rows(data_rows, self.HEADERS)
        assert len(result["customers"]) == 1
        assert result["customers"][0]["fields"]["conclusion"] == "正常"
        assert result["customDimensions"] == []


# ═══════════════════════════════════════════════════════════════════════════════
# 5. D4-31 客户访谈记录 IO 往返
# ═══════════════════════════════════════════════════════════════════════════════


class TestD431Roundtrip:
    """Req 1.4: D4-31 保留问卷结构、q1_relation 多选数组。"""

    HEADERS = _SHEET_HEADERS["D4-31"]

    def test_q1_relation_array_preserved(self):
        """q1_relation 必须是 string[]，不能变成字符串。"""
        data = {
            "target": "ABC公司", "timePlace": "2025-06-15 北京",
            "q1_relation": ["客户是终端客户", "客户是供应商"],
            "q3a_hasContract": "是", "q5_otherMatters": "无",
        }
        from openpyxl import Workbook
        wb = Workbook(); ws = wb.active
        ws.append(self.HEADERS)
        ws.append(_export_d4_31_row(data, self.HEADERS))
        data_rows = list(ws.iter_rows(min_row=2, values_only=True))
        result = _parse_d4_31_row(data_rows[0], self.HEADERS)
        assert result["target"] == "ABC公司"
        assert isinstance(result["q1_relation"], list)
        assert "客户是终端客户" in result["q1_relation"]
        assert "客户是供应商" in result["q1_relation"]
        assert result["q3a_hasContract"] == "是"

    def test_empty_q1_relation(self):
        data = {"target": "DEF公司", "q1_relation": []}
        exported = _export_d4_31_row(data, self.HEADERS)
        result = _parse_d4_31_row(tuple(exported), self.HEADERS)
        assert result["q1_relation"] == []

    def test_four_sections_roundtrip(self):
        """四章节字段全覆盖。"""
        data = {
            "target": "GHI", "timePlace": "上海", "interviewee": "李四", "interviewer": "审计员A",
            "introduction": "受访人介绍",
            "companyName": "GHI公司", "regCapital": "1000万", "establishDate": "2020-01",
            "bizNature": "有限公司", "legalRep": "王五", "equityStructure": "自然人100%",
            "q1_relation": ["客户是经销商(客户)"], "q2a_payment": "赊销方式", "q2b_collection": "货到收款方式",
            "q6_hasShares": "否", "q6_hasPosition": "否", "q6_hasTransaction": "否",
            "signInterviewee": "李四", "signAuditor": "审计员A", "signOther": "", "signDate": "2025-06-15",
        }
        from openpyxl import Workbook
        wb = Workbook(); ws = wb.active
        ws.append(self.HEADERS)
        ws.append(_export_d4_31_row(data, self.HEADERS))
        data_rows = list(ws.iter_rows(min_row=2, values_only=True))
        result = _parse_d4_31_row(data_rows[0], self.HEADERS)
        assert result["companyName"] == "GHI公司"
        assert result["q6_hasShares"] == "否"
        assert result["signDate"] == "2025-06-15"


# ═══════════════════════════════════════════════════════════════════════════════
# 6. D4-32 资金流水检查 IO 往返
# ═══════════════════════════════════════════════════════════════════════════════


class TestD432Roundtrip:
    """Req 1.5: D4-32 六组归属由明确分组列恢复，不硬切。"""

    HEADERS = _SHEET_HEADERS["D4-32"]

    def _make_group(self, key: str, rows: list[dict]) -> dict:
        return {"key": key, "rows": rows}

    def _make_row(self, name: str, amount=0, anomaly="否") -> dict:
        return {
            "id": f"ff-{uuid4().hex[:12]}", "name": name, "amount": amount,
            "ratio": "", "bank": "", "account": "", "method": "", "hasAnomaly": anomaly, "indexRef": "",
        }

    def test_roundtrip_six_groups(self):
        groups = [
            self._make_group("supplier", [self._make_row("供应商A", 100000)]),
            self._make_group("customer", [self._make_row("客户B", 200000, "是")]),
            self._make_group("shareholder", []),
            self._make_group("controller", [self._make_row("控制人C", 50000)]),
            self._make_group("management", []),
            self._make_group("related", [self._make_row("关联方D", 30000)]),
        ]
        from openpyxl import Workbook
        wb = Workbook(); ws = wb.active
        ws.append(self.HEADERS)
        _export_d4_32_rows(ws, groups, self.HEADERS)
        data_rows = list(ws.iter_rows(min_row=2, values_only=True))
        result = _parse_d4_32_rows(data_rows, self.HEADERS)
        assert len(result) == 6
        by_key = {g["key"]: g for g in result}
        assert len(by_key["supplier"]["rows"]) == 1
        assert by_key["supplier"]["rows"][0]["name"] == "供应商A"
        assert len(by_key["customer"]["rows"]) == 1
        assert by_key["customer"]["rows"][0]["hasAnomaly"] == "是"
        assert len(by_key["shareholder"]["rows"]) == 0
        assert len(by_key["controller"]["rows"]) == 1
        assert by_key["controller"]["rows"][0]["name"] == "控制人C"
        assert len(by_key["related"]["rows"]) == 1

    def test_group_order_preserved(self):
        groups = [{"key": k, "rows": []} for k in _D4_32_GROUP_KEYS]
        from openpyxl import Workbook
        wb = Workbook(); ws = wb.active
        ws.append(self.HEADERS)
        _export_d4_32_rows(ws, groups, self.HEADERS)
        data_rows = list(ws.iter_rows(min_row=2, values_only=True))
        result = _parse_d4_32_rows(data_rows, self.HEADERS)
        assert [g["key"] for g in result] == _D4_32_GROUP_KEYS

    def test_unknown_group_no_guess(self):
        """未知分组名不应自动归「其他」，应留在当前组。"""
        # 如果出现非六组标签的行名，当作数据行保留在当前组
        rows = [
            ("", "主要供应商", None, None, None, None, None, None, None),
            (1, "供应商X", 100, "10%", "工行", "123", "直接获取", "否", ""),
            ("", "未知分类", None, None, None, None, None, None, None),  # 不是标准六组
            (2, "对象Y", 200, "20%", "建行", "456", "间接", "是", ""),
        ]
        result = _parse_d4_32_rows(rows, self.HEADERS)
        by_key = {g["key"]: g for g in result}
        # 未知分类不是标准组标签，所以"未知分类"被当作普通行、"对象Y"紧跟在 supplier 组
        assert len(by_key["supplier"]["rows"]) >= 2  # 供应商X + 对象Y（可能含未知分类行）


# ═══════════════════════════════════════════════════════════════════════════════
# 7. 变异检验锚点
# ═══════════════════════════════════════════════════════════════════════════════


class TestMutationAnchors:
    """变异检验：改一处映射必须让至少一条测试变红。"""

    def test_mutation_d4_29_key_change(self):
        """如果 creditCode 映射被改，往返必不一致。"""
        assert _D4_29_HEADER_TO_KEY["统一社会信用代码"] == "creditCode"

    def test_mutation_d4_31_q1_is_list(self):
        """q1_relation 必须解析为 list。"""
        row = tuple([""] * len(_SHEET_HEADERS["D4-31"]))
        # 找到 q1_relation 的位置并填入值
        headers = _SHEET_HEADERS["D4-31"]
        idx = headers.index("客户业务关系")
        row_list = list(row)
        row_list[idx] = "终端客户/供应商"
        result = _parse_d4_31_row(tuple(row_list), headers)
        assert isinstance(result["q1_relation"], list)
        assert len(result["q1_relation"]) == 2

    def test_mutation_d4_32_group_label(self):
        """六组标签必须与前端锁死。"""
        assert _D4_32_GROUP_LABELS["supplier"] == "主要供应商"
        assert _D4_32_GROUP_LABELS["related"] == "其他关联方"
        assert len(_D4_32_GROUP_KEYS) == 6

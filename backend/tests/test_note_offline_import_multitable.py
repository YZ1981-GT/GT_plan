"""离线导入每 sheet 多表解析 + 并回生成链结构 + 真实落库（spec Phase 8 Task 9）。

覆盖：
- extract_sheet_tables：一个 sheet 多张表全部解析（旧实现只读第一张）
- merge_imported_tables_into_section：按表序并回、只改 values、保留 row_type/is_total/_cell_meta/绑定、首表镜像
- 导出→导入 round-trip：多表章节导出后再抽取，表数 / 行标签一致
- apply_import OVERWRITE 写回 _tables（不是扁平单表），回执 tables_written 如实
"""
from __future__ import annotations

import sys
from pathlib import Path

import pytest

_BACKEND = Path(__file__).resolve().parents[1]
if str(_BACKEND) not in sys.path:
    sys.path.insert(0, str(_BACKEND))

from app.services.note_offline_import_tables import (  # noqa: E402
    OfflineImportProtectedCellError,
    OfflineImportStructureError,
    count_changed_tables,
    extract_sheet_tables,
    merge_imported_tables_into_section,
    merge_selected_cells_into_section,
)


def test_structure_driven_two_tables_with_in_table_blank_rows() -> None:
    """结构驱动：表内有空明细行也不会被误切（真实验收 2→4 的根因回归）。"""
    grid = [
        ["其他权益工具投资情况"],
        ["项目", "期末余额", "期初余额"],
        [None, None, None],   # 表内空明细行
        [None, None, None],
        ["合计", 0, 0],
        [None, None, None],   # 表间空行
        ["期末其他权益工具投资情况"],
        ["项目名称", "本期确认的股利收入"],
        [None, None],
        ["合计", 0],
    ]
    expected = [
        {"name": "其他权益工具投资情况", "headers": ["项目", "期末余额", "期初余额"],
         "rows": [{"label": "", "values": [None, None]}, {"label": "", "values": [None, None]},
                  {"label": "合计", "values": [0, 0]}]},
        {"name": "期末其他权益工具投资情况", "headers": ["项目名称", "本期确认的股利收入"],
         "rows": [{"label": "", "values": [None]}, {"label": "合计", "values": [0]}]},
    ]
    tables = extract_sheet_tables(grid, expected=expected)
    assert len(tables) == 2, f"结构驱动应仍为 2 张表（表内空行不切），实为 {len(tables)}"
    assert len(tables[0]["rows"]) == 3 and len(tables[1]["rows"]) == 2
    assert tables[0]["rows"][2]["cells"] == ["合计", 0, 0]
    assert tables[1]["rows"][1]["cells"] == ["合计", 0]


def test_fallback_two_tables_in_one_sheet() -> None:
    """无既有结构（IMPORT_ONLY）退回空行分块：表A 空行 表B。"""
    grid = [
        ["表A"],
        ["项目", "期末", "期初"],
        ["货币资金", 100, 90],
        ["银行存款", 80, 70],
        [None, None, None],
        ["表B"],
        ["项目", "本期"],
        ["销售费用", 50],
    ]
    tables = extract_sheet_tables(grid)
    assert len(tables) == 2, f"应解析出 2 张表，实为 {len(tables)}"
    assert tables[0]["name"] == "表A" and tables[1]["name"] == "表B"
    assert [r["cells"][0] for r in tables[0]["rows"]] == ["货币资金", "银行存款"]
    assert [r["cells"][0] for r in tables[1]["rows"]] == ["销售费用"]


def test_extract_tables_split_by_blank_row() -> None:
    """两表以空行分隔（导出固定在表间空一行）；块内第一行是表头、其余是数据行。"""
    grid = [
        ["项目", "期末"],
        ["A", 1],
        ["合计", 1],
        [None, None],
        ["项目", "本期"],
        ["B", 2],
    ]
    tables = extract_sheet_tables(grid)
    assert len(tables) == 2
    assert [r["cells"][0] for r in tables[0]["rows"]] == ["A", "合计"]
    assert [r["cells"][0] for r in tables[1]["rows"]] == ["B"]


def test_extract_ignores_single_column_only_block() -> None:
    """纯单列块（章首指引，无 value 列头）不当作表。"""
    grid = [["【提示：本章仅在…时填列】"], [None, None], ["项目", "期末"], ["A", 1]]
    tables = extract_sheet_tables(grid)
    assert len(tables) == 1 and [r["cells"][0] for r in tables[0]["rows"]] == ["A"]


def test_merge_updates_values_and_preserves_non_value_fields() -> None:
    existing = {
        "_tables": [
            {
                "name": "货币资金", "headers": ["项目", "期末", "期初"],
                "rows": [
                    {"label": "库存现金", "values": [0, 0], "row_type": "data",
                     "_cell_meta": {"0": {"binding_id": "八、1.库存现金.end"}}},
                    {"label": "合计", "values": [None, None], "is_total": True},
                ],
            },
            {"name": "受限", "headers": ["项目", "金额"], "rows": [{"label": "质押", "values": [0]}]},
        ],
        "headers": ["项目", "期末", "期初"], "rows": [{"label": "库存现金", "values": [0, 0]}],
    }
    imported = [
        {"name": "货币资金", "headers": ["项目", "期末", "期初"],
         "rows": [{"cells": ["库存现金", 0, 90]}, {"cells": ["合计", 100, 90]}]},
        {"name": "受限", "headers": ["项目", "金额"], "rows": [{"cells": ["质押", 50]}]},
    ]
    out = merge_imported_tables_into_section(existing, imported)
    t0 = out["_tables"][0]
    assert t0["rows"][0]["values"] == [0, 90], "未保护列的值未并入"
    assert t0["rows"][0]["row_type"] == "data", "row_type 丢失"
    assert t0["rows"][0]["_cell_meta"]["0"]["binding_id"] == "八、1.库存现金.end", "绑定丢失"
    assert t0["rows"][1]["is_total"] is True, "is_total 丢失"
    assert out["_tables"][1]["rows"][0]["values"] == [50], "第二张表未并入"
    # 首表镜像
    assert out["rows"] == t0["rows"] and out["headers"] == t0["headers"]
    # 不改入参
    assert existing["_tables"][0]["rows"][0]["values"] == [0, 0]


def test_merge_single_table_stays_flat() -> None:
    existing = {"headers": ["项目", "期末"], "rows": [{"label": "A", "values": [0]}]}
    imported = [{"name": "", "headers": ["项目", "期末"], "rows": [{"cells": ["A", 7]}]}]
    out = merge_imported_tables_into_section(existing, imported)
    assert "_tables" not in out, "单表章节不应平白升级为 _tables"
    assert out["rows"][0]["values"] == [7]


def test_merge_import_only_builds_structure_from_scratch() -> None:
    out = merge_imported_tables_into_section({}, [
        {"name": "表1", "headers": ["项目", "本期"], "rows": [{"cells": ["X", 3]}]},
        {"name": "表2", "headers": ["项目", "本期"], "rows": [{"cells": ["Y", 4]}]},
    ])
    assert len(out["_tables"]) == 2
    assert out["_tables"][0]["rows"][0] == {"label": "X", "values": [3]}
    assert out["rows"] == out["_tables"][0]["rows"]


def test_roundtrip_export_then_extract_all_tables() -> None:
    """多表章节导出 → 再抽取，表数与各表行标签一致（证明 export↔extract 口径对齐）。"""
    from openpyxl import load_workbook
    from io import BytesIO

    from app.services.note_offline_export_service import export_sections_to_xlsx

    section = {
        "section_id": "八、1", "note_section": "八、1", "section_title": "货币资金",
        "table_data": {
            "_tables": [
                {"name": "货币资金", "headers": ["项目", "期末", "期初"],
                 "rows": [{"label": "库存现金", "values": [100, 90]}, {"label": "银行存款", "values": [80, 70]}]},
                {"name": "受限货币资金", "headers": ["项目", "金额"],
                 "rows": [{"label": "保证金", "values": [5]}]},
            ],
        },
    }
    xlsx_bytes, _names = export_sections_to_xlsx([section], project_name="测试", include_formulas=False)
    # 定位该 section 的 sheet：第 1 行标题 / 第 2 行 section_id
    wb = load_workbook(BytesIO(xlsx_bytes), data_only=True)
    target = None
    for ws in wb.worksheets:
        c = ws.cell(row=2, column=1).value
        if c and str(c).startswith("section_id:八、1"):
            target = ws
            break
    assert target is not None, "导出包里找不到八、1 的 sheet"
    grid = [[target.cell(row=r, column=col).value for col in range(1, target.max_column + 1)]
            for r in range(3, target.max_row + 1)]
    # 结构驱动：传既有 _tables 作锚点（真实路径即如此，表内空行不会被误切）
    tables = extract_sheet_tables(grid, expected=section["table_data"]["_tables"])
    assert len(tables) == 2, f"round-trip 后应仍为 2 张表，实为 {len(tables)}"
    assert [r["cells"][0] for r in tables[0]["rows"]] == ["库存现金", "银行存款"]
    assert [r["cells"][0] for r in tables[1]["rows"]] == ["保证金"]


def test_apply_import_writes_tables_not_flat() -> None:
    """apply_import OVERWRITE 后 table_data 带 _tables（多表），tables_written 回执如实。"""
    from io import BytesIO

    from openpyxl import load_workbook  # noqa: F401 (确保依赖在)

    from app.services.note_offline_export_service import export_sections_to_xlsx
    from app.services.note_offline_import_service import apply_import, ConflictResolution

    section = {
        "section_id": "八、1", "note_section": "八、1", "section_title": "货币资金",
        "table_data": {
            "_tables": [
                {"name": "货币资金", "headers": ["项目", "期末", "期初"],
                 "rows": [{"label": "库存现金", "values": [0, 0]}]},
                {"name": "受限货币资金", "headers": ["项目", "金额"],
                 "rows": [{"label": "保证金", "values": [0]}]},
            ],
        },
    }
    import copy

    imported = copy.deepcopy(section)
    imported["table_data"]["_tables"][0]["rows"][0]["values"] = [100, 90]
    imported["table_data"]["_tables"][1]["rows"][0]["values"] = [5]
    xlsx_bytes, _ = export_sections_to_xlsx([imported], project_name="测试", include_formulas=False)
    result = apply_import(xlsx_bytes, [section], {"八、1": ConflictResolution.OVERWRITE})
    assert result.sections_imported == 1
    assert result.tables_written == 2, f"回执应报 2 张表写入，实为 {result.tables_written}"
    td = result.applied_sections["八、1"]
    assert isinstance(td.get("_tables"), list) and len(td["_tables"]) == 2, "未写回多表结构"


def test_headers_without_label_column_roundtrip() -> None:
    """模板 headers 只含数据列时，导出会补「项目」；结构解析必须按同一规则识别。"""
    grid = [["项目", "期末"], ["A", 1]]
    expected = [{"headers": ["期末"], "rows": [{"label": "A", "values": [0]}]}]
    assert extract_sheet_tables(grid, expected=expected)[0]["rows"][0]["cells"] == ["A", 1]


def test_structure_change_is_rejected_instead_of_silently_dropping_rows() -> None:
    import pytest

    expected = [{"name": "明细", "headers": ["项目", "金额"],
                 "rows": [{"label": "A", "values": [0]}]}]
    with pytest.raises(OfflineImportStructureError, match="只允许改值"):
        extract_sheet_tables([["项目", "金额"], ["A", 1], ["★ 新增", 2]], expected=expected)


def test_row_reorder_is_rejected_instead_of_positional_overwrite() -> None:
    import pytest

    existing = {"headers": ["项目", "金额"], "rows": [
        {"label": "A", "values": [1]}, {"label": "B", "values": [2]},
    ]}
    imported = [{"headers": ["项目", "金额"], "rows": [
        {"cells": ["B", 20]}, {"cells": ["A", 10]},
    ]}]
    with pytest.raises(OfflineImportStructureError, match="行结构已变化"):
        merge_imported_tables_into_section(existing, imported)


def test_merge_selected_cells_supports_second_table_and_keeps_other_values() -> None:
    existing = {"_tables": [
        {"name": "表1", "headers": ["项目", "金额"], "rows": [{"label": "A", "values": [1]}]},
        {"name": "表2", "headers": ["项目", "金额", "期初"],
         "rows": [{"label": "B", "values": [2, 3], "row_type": "data"}]},
    ]}
    imported = [
        {"name": "表1", "headers": ["项目", "金额"], "rows": [{"cells": ["A", 10]}]},
        {"name": "表2", "headers": ["项目", "金额", "期初"], "rows": [{"cells": ["B", 20, 30]}]},
    ]
    out = merge_selected_cells_into_section(existing, imported, ["1:0:1"])
    assert out["_tables"][0]["rows"][0]["values"] == [1]
    assert out["_tables"][1]["rows"][0]["values"] == [2, 30]
    assert out["_tables"][1]["rows"][0]["row_type"] == "data"
    assert out["rows"] == out["_tables"][0]["rows"]
    assert count_changed_tables(existing, out) == 1


def test_apply_import_rejects_import_only_and_noop_receipt_is_zero() -> None:
    from app.services.note_offline_export_service import export_sections_to_xlsx
    from app.services.note_offline_import_service import ConflictResolution, apply_import

    current = {"section_id": "s1", "section_title": "现有", "table_data": {
        "headers": ["项目", "金额"], "rows": [{"label": "A", "values": [1]}],
    }}
    new = {"section_id": "s2", "section_title": "新增", "table_data": {
        "headers": ["项目", "金额"], "rows": [{"label": "A", "values": [2]}],
    }}
    same_bytes, _ = export_sections_to_xlsx([current])
    same = apply_import(same_bytes, [current], {"s1": ConflictResolution.OVERWRITE})
    assert same.success is True and same.sections_imported == 0 and same.tables_written == 0

    new_bytes, _ = export_sections_to_xlsx([new])
    rejected = apply_import(new_bytes, [current], {"s2": ConflictResolution.OVERWRITE})
    assert rejected.success is False and rejected.sections_imported == 0
    assert rejected.tables_written == 0 and "不创建新章节" in rejected.errors[0]


def test_missing_expected_table_header_is_rejected_not_silently_kept() -> None:
    import pytest

    expected = [{"name": "明细", "headers": ["项目", "金额"], "rows": [{"label": "A", "values": [1]}]}]
    with pytest.raises(OfflineImportStructureError, match="找不到既有表"):
        extract_sheet_tables([["项目", "被改名"], ["A", 2]], expected=expected)


def _protected_table_data(marker: str) -> dict:
    row = {"label": "A", "values": [1, 2]}
    root = {"headers": ["项目", "金额", "期初"], "rows": [row]}
    if marker == "binding":
        row["_cell_meta"] = {"0": {"binding_id": "bind-1"}}
    elif marker in {"auto", "formula", "locked"}:
        row["_cell_modes"] = {"0": marker}
    elif marker == "source":
        row["_cell_meta"] = {"0": {"source": "trial_balance"}}
    elif marker == "legacy_binding":
        root["_bindings"] = {"0:0": {"source": "trial_balance"}}
    elif marker == "legacy_formula":
        root["_formulas"] = {"0:0": {"expression": "=1+1"}}
    return root


@pytest.mark.parametrize(
    "marker", ["binding", "auto", "formula", "locked", "source", "legacy_binding", "legacy_formula"],
)
def test_overwrite_rejects_changed_protected_cells(marker: str) -> None:
    existing = _protected_table_data(marker)
    imported = [{"headers": ["项目", "金额", "期初"], "rows": [{"cells": ["A", 9, 2]}]}]
    with pytest.raises(OfflineImportProtectedCellError, match="不得改写"):
        merge_imported_tables_into_section(existing, imported)


def test_protected_equal_value_is_allowed_and_manual_value_can_change() -> None:
    protected = _protected_table_data("binding")
    same = [{"headers": ["项目", "金额", "期初"], "rows": [{"cells": ["A", 1.0, 2]}]}]
    assert merge_imported_tables_into_section(protected, same)["rows"][0]["values"] == [1.0, 2]

    manual = _protected_table_data("manual")
    changed = [{"headers": ["项目", "金额", "期初"], "rows": [{"cells": ["A", 9, 2]}]}]
    assert merge_imported_tables_into_section(manual, changed)["rows"][0]["values"] == [9, 2]


def test_merge_checks_only_selected_cells_against_current_db_protection() -> None:
    existing = _protected_table_data("binding")
    imported = [{"headers": ["项目", "金额", "期初"], "rows": [{"cells": ["A", 9, 20]}]}]
    out = merge_selected_cells_into_section(existing, imported, ["0:1"])
    assert out["rows"][0]["values"] == [1, 20], "未选择的受保护变化不得阻断手工格合并"
    with pytest.raises(OfflineImportProtectedCellError, match="binding"):
        merge_selected_cells_into_section(existing, imported, ["0:0"])


def test_merge_rejects_protected_cell_in_second_table() -> None:
    existing = {"_tables": [
        {"name": "表1", "headers": ["项目", "金额"], "rows": [{"label": "A", "values": [1]}]},
        {"name": "表2", "headers": ["项目", "金额"], "rows": [
            {"label": "B", "values": [2], "_cell_modes": {"0": "locked"}},
        ]},
    ]}
    imported = [
        {"name": "表1", "headers": ["项目", "金额"], "rows": [{"cells": ["A", 1]}]},
        {"name": "表2", "headers": ["项目", "金额"], "rows": [{"cells": ["B", 3]}]},
    ]
    with pytest.raises(OfflineImportProtectedCellError, match="1:0:0"):
        merge_selected_cells_into_section(existing, imported, ["1:0:0"])

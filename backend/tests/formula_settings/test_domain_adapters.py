"""底稿 / 附注 / 合并附注三个域的适配器与消费方接线（tasks 6.1~6.4，design §12）。

SQLite 真 ORM；每条用例都经真实存储（``wp_formula`` / ``formula_setting`` / ``consol_note_formula``）落库再读回。
"""
from __future__ import annotations

import uuid

import openpyxl
import pytest
import sqlalchemy as sa

from app.models.consol_push_models import ConsolNoteFormula
from app.models.formula_setting_models import FormulaSetting
from app.models.report_models import DisclosureNote
from app.models.workpaper_models import WorkingPaper, WpFormula, WpIndex, WpSourceType
from app.services.formula_settings import note_settings, propagation
from app.services.formula_settings.errors import FormulaSettingError, VersionConflict, VersionRequired
from app.services.formula_settings.project_ops import resolve_items, restore_global
from app.services.formula_settings.publish import publish_global
from app.services.formula_settings.stores import get_store
from app.services.formula_settings.wp_effective import (
    effective_cell_formulas,
    init_prefill_overrides,
    overlay_prefill_formulas,
    override_value,
)
from app.services.prefill_engine import _FORMULA_RE

from ._env import add_project, add_user, make_db, pending_items

EXTRA = [WpIndex, WorkingPaper, WpFormula, DisclosureNote, ConsolNoteFormula]
SLOT = "D2|审定表D2-1|E8"
G1, G2, MINE = "TB('1122','期末余额')", "TB('1122','审定数')", "TB('1122','期末余额')-TB('1231','期末余额')"


async def add_wp(db, project_id, code: str = "D2") -> uuid.UUID:
    idx = WpIndex(id=uuid.uuid4(), project_id=project_id, wp_code=code, wp_name=code)
    db.add(idx)
    await db.flush()
    wp = WorkingPaper(id=uuid.uuid4(), project_id=project_id, wp_index_id=idx.id, file_path=f"/tmp/{code}.xlsx",
                      source_type=WpSourceType.template)
    db.add(wp)
    await db.flush()
    return wp.id


async def wp_rows(db, project_id) -> list[WpFormula]:
    return list((await db.execute(sa.select(WpFormula).where(WpFormula.project_id == project_id))).scalars().all())


def wdef(expr: str) -> dict:
    return {"expression": expr, "formula_type": "auto_calc"}


# ── 底稿 ──────────────────────────────────────────────────────────────────────


async def test_wp_publish_inherit_follow_conflict_then_adopt():
    async with make_db(EXTRA) as factory, factory() as db:
        admin = await add_user(db, "adm", "admin")
        inherit, follow, conflict = [await add_project(db, n) for n in ("甲", "乙", "丙")]
        await add_project(db, "丁")  # 没有该底稿
        for p in (inherit, follow, conflict):
            await add_wp(db, p)
        await publish_global(db, domain="wp", scope="*", user_id=admin.id,
                             changes=[{"slot_key": SLOT, "definition": wdef(G1)}])
        store = get_store("wp")
        await store.write_project(db, follow, SLOT, definition=wdef(G1), base_definition=wdef(G1), base_version=1,
                                  pinned=False, base_origin="global", base_template_id=None, user_id=admin.id)
        await store.write_project(db, conflict, SLOT, definition=wdef(MINE), base_definition=None, base_version=None,
                                  pinned=False, base_origin=None, base_template_id=None, user_id=admin.id)

        out = await publish_global(db, domain="wp", scope="*", user_id=admin.id,
                                   changes=[{"slot_key": SLOT, "definition": wdef(G2), "expected_version": 1}])
        counts = out["summary"]["counts"]
        assert (counts["inherit"], counts["follow"], counts["conflict"]) == (2, 1, 1)
        assert await wp_rows(db, inherit) == [], "继承：不写本项目行"
        (row,) = await wp_rows(db, follow)
        assert (row.expression, row.formula_source, row.base_version, row.definition_version) == (G2, "preset", 2, 2)
        assert row.base_definition["expression"] == G2 and row.stable_sheet_key and row.definition_hash
        (mine,) = await wp_rows(db, conflict)
        assert mine.expression == MINE, "冲突不覆盖本项目公式"
        (item,) = await pending_items(db, conflict)
        await resolve_items(db, project_id=conflict, user_id=admin.id,
                            resolutions=[{"item_id": str(item.id), "choice": "adopt"}])
        (mine,) = await wp_rows(db, conflict)
        assert (mine.expression, mine.formula_source, mine.base_version) == (G2, "preset", 2)


async def test_wp_store_guards_and_restore():
    async with make_db(EXTRA) as factory, factory() as db:
        user = await add_user(db, "u")
        bare, with_wp = await add_project(db, "无底稿"), await add_project(db, "有底稿")
        await add_wp(db, with_wp)
        store = get_store("wp")
        kw = dict(base_definition=None, base_version=None, pinned=False, base_origin=None, base_template_id=None,
                  user_id=user.id)
        await store.write_project(db, bare, SLOT, definition=wdef(G1), **kw)
        assert await wp_rows(db, bare) == [], "项目没有该底稿：无处可写"
        await store.write_project(db, with_wp, SLOT, definition={"disabled": True}, **kw)
        assert await wp_rows(db, with_wp) == [], "底稿域没有停用：不写"
        for bad in ("eval('1')", "TB('1122'"):
            with pytest.raises(FormulaSettingError):
                await store.write_project(db, with_wp, SLOT, definition=wdef(bad), **kw)
        with pytest.raises(FormulaSettingError):
            await store.write_project(db, with_wp, SLOT, definition={**wdef(G1), "formula_type": "x"}, **kw)
        await store.write_project(db, with_wp, SLOT, definition=wdef("=TB_SUM('1121~1122','期末余额')"), **kw)
        (row,) = await wp_rows(db, with_wp)
        assert row.expression == "TB_SUM('1121~1122','期末余额')", "预设自带的 TB_SUM 可写；前导「=」去掉"
        out = await restore_global(db, domain="wp", project_id=with_wp, slot_keys=[SLOT], user_id=user.id)
        assert out["restored"] == [SLOT] and await wp_rows(db, with_wp) == []


def _workbook():
    wb = openpyxl.Workbook()
    ws = wb.active
    ws.title = "审定表D2-1"
    ws["E8"] = "=TB('1122','期初余额')"   # 模板内置公式：有效公式要替换它
    ws["F8"] = "=TB('1001','期末余额')"   # 不相干的内置公式保留
    ws["C5"] = "未审数"                   # 语义位置：列头
    return wb


async def test_wp_effective_and_prefill_overlay():
    async with make_db(EXTRA) as factory, factory() as db:
        admin = await add_user(db, "adm", "admin")
        pid = await add_project(db, "甲")
        wp_id = await add_wp(db, pid)
        await publish_global(db, domain="wp", scope="*", user_id=admin.id, changes=[
            {"slot_key": SLOT, "definition": wdef(G2)},
            {"slot_key": "D2|审定表D2-1|未审数", "definition": wdef("TB('1122','未审数')")},
            {"slot_key": "D2|审定表D2-1|G9", "definition": wdef(MINE)},
            {"slot_key": "D2|审定表D2-1|H9", "definition": {**wdef("TB('1122','审定数')"), "formula_type": "logic_check"}},
            {"slot_key": "D9|别的底稿|A1", "definition": wdef(G1)},
        ])
        code, cells = await effective_cell_formulas(db, pid, wp_id)
        assert code == "D2" and {(c.cell, c.layer) for c in cells} == {
            ("E8", "global"), ("未审数", "global"), ("G9", "global"), ("H9", "global")}, "只取本底稿编码"
        await get_store("wp").write_project(db, pid, "D2|审定表D2-1|G9", definition=wdef("TB('1231','期末余额')"),
                                            base_definition=None, base_version=None, pinned=False, base_origin=None,
                                            base_template_id=None, user_id=admin.id)
        wb = _workbook()
        scanned = [{"sheet": "审定表D2-1", "cell": c, "original": wb.active[c].value} for c in ("E8", "F8")]
        found, notes = await overlay_prefill_formulas(db, pid, await db.get(WorkingPaper, wp_id), wb, scanned,
                                                      _FORMULA_RE)
        by_cell = {f["cell"]: f for f in found}
        assert by_cell["E8"]["original"] == f"={G2}" and by_cell["E8"]["layer"] == "global", "全局覆盖模板内置公式"
        assert by_cell["F8"]["original"] == "=TB('1001','期末余额')", "无关单元格不动"
        assert by_cell["C7"]["original"] == "=TB('1122','未审数')", "语义位置按初始化同一套规则定位"
        assert by_cell["G9"]["layer"] == "project", "本项目 > 全局"
        assert "H9" not in by_cell and notes == [], "检查类公式不改值"

        await publish_global(db, domain="wp", scope="*", user_id=admin.id, changes=[
            {"slot_key": SLOT, "definition": wdef(MINE), "expected_version": 1}])
        found, notes = await overlay_prefill_formulas(db, pid, await db.get(WorkingPaper, wp_id), _workbook(),
                                                      scanned, _FORMULA_RE)
        assert "E8" not in {f["cell"] for f in found}, "复合公式：不再按模板内置公式算"
        assert [n["cell"] for n in notes] == ["审定表D2-1!E8"] and "公式运行时" in notes[0]["error"]


async def test_wp_deleted_global_falls_back_to_preset_and_init_split():
    preset_slot = "E1|货币资金审定表E1-1|未审数"   # prefill_formula_mapping.json 里的语义单元格
    async with make_db(EXTRA) as factory, factory() as db:
        admin = await add_user(db, "adm", "admin")
        pid = await add_project(db, "甲")
        wp_id = await add_wp(db, pid, "E1")
        await publish_global(db, domain="wp", scope="*", user_id=admin.id, changes=[
            {"slot_key": preset_slot, "definition": wdef("TB('1002','期末余额')")},
            {"slot_key": "E1|货币资金审定表E1-1|K20", "definition": wdef("TB('1012','期末余额')")}])
        user_formulas, overrides = await init_prefill_overrides(db, pid, wp_id)
        assert overrides == {"货币资金审定表E1-1!未审数": "=TB('1002','期末余额')"}, "映射单元格 → 按有效公式取值"
        assert user_formulas == {"货币资金审定表E1-1!K20": {"formula": "=TB('1012','期末余额')", "layer": "global"}}
        await publish_global(db, domain="wp", scope="*", user_id=admin.id, changes=[
            {"slot_key": preset_slot, "definition": None, "expected_version": 1}])
        _, cells = await effective_cell_formulas(db, pid, wp_id)
        (fallback,) = [c for c in cells if c.cell == "未审数"]
        assert fallback.layer == "preset" and "TB('1001','期末余额')" in fallback.expression, \
            "全局删除后按系统预设重算（不残留旧全局公式的值）"


def test_override_value_reads_its_own_accounts():
    tb = {"1001": {"期末余额": 10.0}, "1002": {"期末余额": 20.0}, "1121": {"期初余额": 1.0}, "1122": {"期初余额": 2.0}}
    assert override_value("=TB('1002','期末余额')", "=TB('1001','期末余额')", tb, None, ["1001"]) == 20.0
    assert override_value("TB('1001','期末余额') + TB('1002','期末余额') - 3", "", tb, None, ["1001"]) == 27.0
    assert override_value("TB('1001','期末余额') + TB('9999','期末余额')", "", tb, None, []) == 10.0, "缺的科目按 0"
    assert override_value("TB('9999','期末余额')", "", tb, None, []) is None, "全缺：不写（同预设）"
    assert override_value("=TB_SUM('1121~1122','期初余额')", "=TB_SUM('1121~1122','期末余额')", tb, None,
                          ["1121", "1122"]) == 3.0, "范围同预设：沿用映射科目清单"
    assert override_value("=TB_SUM('1121~1123','期初余额')", "=TB_SUM('1121~1122','期末余额')", tb, None,
                          ["1121", "1122"]) is None, "改了范围：初始化不展开（交重新预填）"
    assert override_value("=ADJ('1001','aje_net')", "", tb, {"1001": {"aje_net": 5}}, ["1001"]) == 5
    assert override_value("=WP('D1','审定表','B5')", "", tb, None, []) is None
    assert override_value("ROUND(TB('1001','期末余额'),2)", "", tb, None, []) is None


def test_prefill_xlsx_uses_mapping_override(tmp_path):
    from app.services.wp_template_init_service import prefill_workpaper_xlsx

    def run(overrides):
        path = tmp_path / f"e1_{len(overrides or {})}.xlsx"
        wb = openpyxl.Workbook()
        ws = wb.active
        ws.title = "货币资金审定表E1-1"
        ws["C5"] = "未审数"
        wb.save(path)
        tb = {"1001": {"期末余额": 10.0}, "1002": {"期末余额": 20.0}, "1012": {"期末余额": 5.0}}
        prefill_workpaper_xlsx(path, "E1", tb, mapping_overrides=overrides)
        return openpyxl.load_workbook(path)["货币资金审定表E1-1"]["C7"].value

    assert run({"货币资金审定表E1-1!未审数": "=TB('1002','期末余额')"}) == 20.0
    assert run({"货币资金审定表E1-1!未审数":
                "=TB('1001','期末余额') + TB('1002','期末余额') + TB('1012','期末余额')"}) == 35.0
    assert run(None) != 20.0, "无覆盖：照原预设口径"


# ── 附注 ──────────────────────────────────────────────────────────────────────

SECTION = "五、1"


async def add_note(db, project_id, year: int = 2025, td: dict | None = None) -> uuid.UUID:
    note = DisclosureNote(id=uuid.uuid4(), project_id=project_id, year=year, note_section=SECTION,
                          section_title="货币资金", table_data=td or {
                              "headers": ["项目", "期末", "期初"], "rows": [],
                              "_formulas": {"2:1": {"type": "vertical_sum", "expression": "SUM_ROWS(0,1)",
                                                    "category": "auto_calc", "source": "template"}}})
    db.add(note)
    await db.flush()
    return note.id


async def _note(db, note_id) -> DisclosureNote:
    return (await db.execute(sa.select(DisclosureNote).where(DisclosureNote.id == note_id)
                             .execution_options(populate_existing=True))).scalar_one()


async def test_note_list_save_materialize_and_cas():
    async with make_db(EXTRA) as factory, factory() as db:
        user = await add_user(db, "u")
        pid = await add_project(db, "甲")
        n25, n24 = await add_note(db, pid), await add_note(db, pid, 2024)
        listing = await note_settings.list_section(db, pid, 2025, SECTION)
        assert [(r["target"], r["layer"]) for r in listing["formulas"]] == [("2:1", "preset")]
        assert listing["loaded"] == {f"{SECTION}|2:1": None}
        unchanged = await note_settings.save_section(db, pid, 2025, SECTION, listing["formulas"],
                                                     loaded=listing["loaded"], user_id=user.id)
        assert unchanged["events"] == [] and (await _note(db, n25)).table_data.get("_user_formulas") is None, \
            "与模板公式相同：不算适配、不物化"
        saved = await note_settings.save_section(db, pid, 2025, SECTION, [
            {"target": "2:1", "formula": "SUM_ROWS(0,1)+1"}, {"target": "3:1", "formula": "ROW(2)"}],
            loaded=listing["loaded"], user_id=user.id)
        assert {(r["target"], r["layer"]) for r in saved["formulas"]} == {("2:1", "project"), ("3:1", "project")}
        assert saved["events"] and saved["loaded"][f"{SECTION}|2:1"] == "L1"
        for nid in (n25, n24):  # 项目层不分年度
            td = (await _note(db, nid)).table_data
            assert [r["target"] for r in td["_user_formulas"]] == ["2:1", "3:1"] and td["_formulas"]
        with pytest.raises(VersionRequired):
            await note_settings.save_section(db, pid, 2025, SECTION, saved["formulas"], loaded=None, user_id=user.id)
        with pytest.raises(VersionConflict):
            await note_settings.save_section(db, pid, 2025, SECTION, saved["formulas"],
                                             loaded={**saved["loaded"], f"{SECTION}|3:1": "L9"}, user_id=user.id)
        with pytest.raises(ValueError):
            await note_settings.save_section(db, pid, 2025, "不存在", [], loaded={}, user_id=user.id)


async def test_note_global_publish_propagates_and_disable():
    async with make_db(EXTRA) as factory:
        async with factory() as db:
            admin = await add_user(db, "adm", "admin")
            pid = await add_project(db, "甲")
            nid = await add_note(db, pid)
            out = await publish_global(db, domain="note", scope="soe", user_id=admin.id, changes=[
                {"slot_key": f"{SECTION}|4:1", "definition": {"formula": "ROW(3)", "category": "auto_calc"}}])
            await db.commit()
        assert out["events"] == [{"project_id": str(pid), "slot_keys": [f"{SECTION}|4:1"]}]
        propagation.set_session_factory(factory)
        try:
            result = await propagation.propagate_note(pid, None, [f"{SECTION}|4:1"])
        finally:
            propagation.set_session_factory(None)
        assert result["materialized"] == 1 and result["marked"] == 1
        async with factory() as db:
            note = await _note(db, nid)
            assert note.is_stale and [(r["target"], r["layer"]) for r in note.table_data["_user_formulas"]] == [
                ("4:1", "global")]
            listing = await note_settings.list_section(db, pid, 2025, SECTION)
            assert listing["loaded"][f"{SECTION}|4:1"] == "G1"
            kept = [r for r in listing["formulas"] if r["target"] != "4:1"]
            await note_settings.save_section(db, pid, 2025, SECTION, kept, loaded=listing["loaded"], user_id=admin.id)
            (row,) = (await db.execute(sa.select(FormulaSetting).where(
                FormulaSetting.domain == "note", FormulaSetting.layer == "project"))).scalars().all()
            assert row.definition == {"disabled": True}, "删掉全局来的公式 = 本项目停用（全局仍在）"
            assert (await _note(db, nid)).table_data["_user_formulas"] == []
            with pytest.raises(FormulaSettingError):
                await publish_global(db, domain="note", scope="soe", user_id=admin.id, changes=[
                    {"slot_key": f"{SECTION}|5:1", "definition": {"formula": " "}}])


async def test_note_legacy_user_formulas_imported_once():
    async with make_db(EXTRA) as factory, factory() as db:
        pid = await add_project(db, "甲")
        nid = await add_note(db, pid, td={
            "headers": [], "rows": [],
            "_formulas": {"2:1": {"type": "vertical_sum", "expression": "SUM_ROWS(0,1)", "category": "auto_calc"}},
            "_user_formulas": [{"target": "2:1", "formula": "SUM_ROWS(0,1)"},     # 同模板公式：不导入
                               {"target": "3:1", "formula": "ROW(2)*2", "category": "logic_check"}]})
        assert await note_settings.materialize(db, pid, [SECTION]) == 1
        rows = (await db.execute(sa.select(FormulaSetting).where(FormulaSetting.domain == "note"))).scalars().all()
        assert [(r.slot_key, r.definition["formula"], r.definition["category"]) for r in rows] == [
            (f"{SECTION}|3:1", "ROW(2)*2", "logic_check")]
        td = (await _note(db, nid)).table_data
        assert td["_user_formulas_origin"] == "formula_settings" and [r["target"] for r in td["_user_formulas"]] == [
            "3:1"]
        assert await note_settings.materialize(db, pid, [SECTION]) == 0, "幂等"


# ── 合并附注 ──────────────────────────────────────────────────────────────────

CSEC, CELL = "五-1-1", "五-1-1|5|1"


async def test_consol_note_effective_and_kernel():
    from app.services.formula_settings.consol_note_settings import effective_formulas

    async with make_db(EXTRA) as factory, factory() as db:
        admin = await add_user(db, "adm", "admin")
        follow, mine = await add_project(db, "乙"), await add_project(db, "丙")
        await publish_global(db, domain="consol_note", scope="soe", user_id=admin.id, changes=[
            {"slot_key": CELL, "definition": {"formula": "TB('1122','期末余额')"}}])
        g = (await db.execute(sa.select(ConsolNoteFormula))).scalar_one()
        assert (g.version, g.source) == (1, "manual")
        store = get_store("consol_note")
        base = {"formula": "TB('1122','期末余额')", "description": None}
        await store.write_project(db, follow, CELL, definition=dict(base), base_definition=base, base_version=1,
                                  pinned=False, base_origin="global", base_template_id=None, user_id=admin.id)
        await store.write_project(db, mine, CELL, definition={"formula": "TB('1122','期末余额')*2"},
                                  base_definition=base, base_version=1, pinned=False, base_origin="global",
                                  base_template_id=None, user_id=admin.id)
        await store.write_project(db, mine, "五-1-1|6|1", definition={"disabled": True, "formula": "ROW(9)"},
                                  base_definition=None,
                                  base_version=None, pinned=False, base_origin=None, base_template_id=None,
                                  user_id=admin.id)
        globals_ = [g, type("G", (), {"row_index": 6, "col_index": 1, "formula": "ROW(1)", "source": "seed"})()]
        eff = await effective_formulas(db, mine, CSEC, globals_)
        assert [(f.row_index, f.formula, f.source) for f in eff] == [(5, "TB('1122','期末余额')*2", "project")], \
            "本项目覆盖全局、停用的单元格剔除"
        assert await effective_formulas(db, follow, "五-9-9", globals_) == globals_, "本章节无适配：原样用全局"

        out = await publish_global(db, domain="consol_note", scope="soe", user_id=admin.id, changes=[
            {"slot_key": CELL, "definition": {"formula": "TB('1122','审定数')"}, "expected_version": 1}])
        assert (out["summary"]["counts"]["follow"], out["summary"]["counts"]["conflict"]) == (1, 1)
        assert g.version == 2 and g.formula == "TB('1122','审定数')"
        assert (await store.list_project(db, follow))[CELL].definition["formula"] == "TB('1122','审定数')"
        assert len(await pending_items(db, mine)) == 1
        with pytest.raises(FormulaSettingError):
            await store.write_project(db, mine, CELL, definition={"formula": "TB('1122'"}, base_definition=None,
                                      base_version=None, pinned=False, base_origin=None, base_template_id=None,
                                      user_id=admin.id)
        with pytest.raises(FormulaSettingError) as exc:
            await publish_global(db, domain="consol_note", scope="soe", user_id=admin.id, changes=[
                {"slot_key": "五-1-1|5|99", "definition": {"formula": "TB('1122','期末余额')"}}])
        assert "列号" in str(exc.value)


async def test_note_materialized_effective_drives_executor_merge():
    """执行层口径：materialize 写进 _user_formulas 的有效定义，经 note_formula_generator._merge_formula_sets
    合并时**用户集优先覆盖模板预设**（执行器正是消费这个合并结果）。证明「改全局公式 → 执行算的值随之变」。"""
    from app.services.note_formula_generator import _merge_formula_sets

    async with make_db(EXTRA) as factory, factory() as db:
        admin = await add_user(db, "adm", "admin")
        pid = await add_project(db, "执行层")
        nid = await add_note(db, pid)  # 模板预设 _formulas 里 2:1 = SUM_ROWS(0,1)
        # 全局发布覆盖同一目标 2:1，并新增 4:1
        await publish_global(db, domain="note", scope="soe", user_id=admin.id, changes=[
            {"slot_key": f"{SECTION}|2:1", "definition": {"formula": "SUM_ROWS(0,1)+99", "category": "auto_calc"}},
            {"slot_key": f"{SECTION}|4:1", "definition": {"formula": "ROW(3)", "category": "auto_calc"}}])
        await db.flush()
        await note_settings.materialize(db, pid, [SECTION])
        await db.flush()
        td = (await _note(db, nid)).table_data
        merged = _merge_formula_sets(td.get("_formulas") or {}, td.get("_user_formulas"))
        # 2:1 用户集（全局有效定义）覆盖模板预设；4:1 是新增；执行器拿到的就是这个集合
        assert merged["2:1"]["expression"] == "SUM_ROWS(0,1)+99"
        assert "4:1" in merged and merged["4:1"]["expression"] == "ROW(3)"

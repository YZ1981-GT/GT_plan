"""test_x3_column_property —— 判据层（常量 / 登记表 / fixture / 纯函数 / 取证 helper）。

从 `backend/tests/test_x3_column_property.py` 拆出：原文件 1166 行 > pre-commit 的 800 行门禁。
**不加 file_size_whitelist** —— 白名单表头写明「仅历史大文件」，新增文件
套用属滥用。

刻意与用例文件**同目录**：判据里的 `Path(__file__).parents[N]` 路径推算
移到子目录会整体错一层（实测过，会变成 fixture setup 全 ERROR）。

用例层的 import 清单由拆分脚本按其**实际引用**算出，不手写 —— 漏一个名字
就是 collection error，会让整份守卫的断言零执行而表面上「没有失败」。

原文件 docstring 原样保留在下方。
"""

"""Feature: x3-adjustment-entry-import-export, Property 3: 导出列面与 sheet 名对齐源模板

*For any* 16 张目标 sheet，``X3_Exporter`` 产出的列头序列与 openpyxl 直读
``backend/wp_templates/`` 对应 tab 第 5 行的标签序列逐字相等（等势且同序），
且导出工作簿的 sheet 名等于 catalog ``sheet_name``。

**Validates: Requirements 3.1, 3.2, 3.6, 3.7, 3.8**

═══════════════════════════════════════════════════════════════════════════════
与 GS2（``test_x3_column_alignment.py``）的分工 —— 本文件**不是**第二份三向比对
═══════════════════════════════════════════════════════════════════════════════

| 面 | GS2（21 passed） | 本文件（Property 3） |
|---|---|---|
| 判据对象 | **常量与登记值**：清单 ``column_map[].col`` ↔ 派生常量 ``COLUMN_ORDER`` ↔ 实读第 5 行 | **运行期产物**：``build_template_workbook`` / ``build_data_workbook`` 产出的 workbook 对象与真实 xlsx 字节 |
| 量化对象 | 16 张 sheet，各断言一次（值比对） | ∀ 任意行数据（行数 / 字段缺失 / 多余字段 / ``None`` / 长中文 / 特殊字符 / 极大极小有限小数）× ∀ 任意列面变形 × ∀ 任意 catalog 替身 |
| 抓的形态 | 三份登记值互相分叉 | 产物列面/列序/宽度/占位位/ sheet 名**因输入或实现改动而变形** |

两处刻意不重复：

* **列面真源可替换性**（``COLUMN_ORDER`` 是否真由清单 ``column_map`` 派生）已由 GS1
  的 ``TestColumnOrderDerivedFromLedger`` 用替身清单驱动覆盖（GD-1 的收口点），本文件
  不再抄一份。
* **工厂 7 列 ``_ADJ_HEADERS`` 那份具体错值**由 GS2 的两条断言承载；本文件承载的是
  更一般的性质 —— **任意**缺列 / 多列 / 换序 / 重复列的候选列面一律被判错（R3.2 / R3.6
  的 ∀ 形态），故不必逐个点名具体错值。

本文件补的是 GD-1 的**同族另一维**（GS1 / GS2 皆无判据）：

    **sheet 名是否真由 ACNR catalog 派生。**

实测清单 ``provenance.column_source.tab`` 与 catalog ``sheet_name`` 对 16 张**逐字相同**
⇒ 把取值改成读清单 tab 是一次**值等价**改动，任何「值比对」判据（含 GS2 的
``test_class_b_sheet_name_taken_from_catalog``）对它恒绿。唯一能抓住的判据形态是
**替身驱动**：换掉 catalog 的 ``sheet_name``，导出产物的 sheet 名必须跟着变（P3-4）。

═══════════════════════════════════════════════════════════════════════════════
期望值的唯一来源 = openpyxl 直读源模板（R3.7）
═══════════════════════════════════════════════════════════════════════════════

本文件**不写任何列头 / sheet 名 / X-3 键字面量**，也不写「共 10 列」这个数字：

* 目标 sheet 集合 —— 从 ``backend/data/adjustment_ie_contract.json`` 按「条目登记了
  ``key_family``」选出（与实现 ``_select_x3_entries`` 同一口径，不另立第二份作业面）。
* 源模板路径 / tab / 表头行号 / 列跨度 —— 取自清单 ``provenance.column_source`` 的
  ``file`` / ``tab`` / ``header_row`` / ``cells``（登记的取值溯源）。**列数由 ``cells``
  的跨度算出**（``A5:J5`` ⇒ 10），不在本文件写死。
* 列头标签本体 —— 每次运行都由 ``load_workbook(read_only=True, data_only=True)`` 现读。
* 导出 sheet 名的期望值 —— ACNR catalog ``sheet_name``（R3.8 真源）。

``backend/wp_templates/`` 只读（R11.1）：只对本文件真正打开的那 16 个文件做
``(size, mtime_ns)`` 快照并前后比对（整目录快照是 GS2 的作业面，不重复）。

═══════════════════════════════════════════════════════════════════════════════
反空转与非退化
═══════════════════════════════════════════════════════════════════════════════

* ``test_anchor_*`` 一组前置锚点：作业面规模 / 列数 / sheet 名非空 / 占位位分布
  ``{1 个: 14 张, 2 个: 1 张, 6 个: 1 张}``。**作业面塌成 0 张或列数塌成 0 时，
  下面所有 ∀ 属性都会恒真** ⇒ 锚点先钉死规模。
* **P3-3 专门证明判据不是恒真**：对同一份真实产物施加 ∀ 变形（缺列 / 多列 / 换序 /
  重复列 / 改 sheet 名 / 跳过占位列 / 整行错位），判据必须报出**对应类别**的违规；
  同一次迭代里先跑未变形的 CONTROL 断言「零违规」，使「变形后打红」可归因。
* 占位位一律由 ``spec.field_keys`` 的 ``None`` 位判定，**不假定只有 F 列**
  （实测 14 张 1 个 · ``L2-3`` 2 个 · ``N5-3`` 6 个）。
"""

from __future__ import annotations

import importlib.util
import io
import json
import re
import string
import sys
import types
from collections.abc import Mapping, Sequence
from dataclasses import dataclass
from functools import lru_cache
from pathlib import Path
from typing import Any, NamedTuple
from unittest import mock

import pytest
from hypothesis import HealthCheck, given, settings
from hypothesis import strategies as st
from openpyxl import load_workbook

from app.routers.wp_render_strategies import _x3_adjustment_import_export as impl

_REPO = Path(__file__).resolve().parents[2]
_TEMPLATES_DIR = _REPO / "backend" / "wp_templates"
_CONTRACT = _REPO / "backend" / "data" / "adjustment_ie_contract.json"

#: 被测出口模块（本任务禁改）。替身重 exec 用它的磁盘路径。
_IMPL_NAME = impl.__name__
_IMPL_PATH = Path(str(impl.__file__)).resolve()

#: 每条属性的迭代次数（任务要求 ≥100；design §Testing Strategy 对 10 条 Property 同）。
_RUNS = 150
#: 真实 xlsx 字节往返那条每次要 save + load 16 次，单例成本高一档，仍 ≥100。
_RUNS_ROUNDTRIP = 100
#: 替身重 exec 那条每次要重新装载一份模块副本，仍 ≥100。
_RUNS_STANDIN = 120

#: 合成标签 / 合成 sheet 名后缀的探针记号 —— 真实列面与真实 sheet 名里都不会出现。
_MARK = "‡属性探针"

_SETTINGS = settings(
    max_examples=_RUNS,
    deadline=None,
    suppress_health_check=[HealthCheck.too_slow],
)


# ═══════════════════════════════════════════════════════════════════════════
# 1. 真源装载（清单 → 源模板实读 → catalog）
# ═══════════════════════════════════════════════════════════════════════════


class _Live(NamedTuple):
    """一张 X-3 的源模板实读结果（期望值侧，唯一取值处）。"""

    sheet_code: str
    path: Path
    tab: str
    header_row: int
    width: int
    labels: tuple[str, ...]
    #: 源模板示例行（表头行下一行）。仅 ``L2-3`` 有实质内容（design E13）。
    sample_row: tuple[str, ...]


@dataclass(frozen=True)
class _LiveBundle:
    templates: dict[str, _Live]
    before: dict[str, tuple[int, int]]
    after: dict[str, tuple[int, int]]


_CELLS_RE = re.compile(r"^([A-Z]{1,2})(\d+):([A-Z]{1,2})(\d+)$")


def _col_index(letters: str) -> int:
    """列字母 → 1-based 序号（``A`` ⇒ 1、``J`` ⇒ 10）。"""
    value = 0
    for ch in letters:
        value = value * 26 + (ord(ch) - ord("A") + 1)
    return value


def _span_width(cells: str, header_row: int, where: str) -> int:
    """从登记的 ``cells``（如 ``A5:J5``）算出列数 —— 列数不在本文件写死。"""
    match = _CELLS_RE.match(cells)
    assert match, f"{where}: `provenance.column_source.cells` 形态非法: {cells!r}"
    left, row_a, right, row_b = match.groups()
    assert row_a == row_b == str(header_row), (
        f"{where}: `cells`={cells!r} 的行号与 `header_row`={header_row} 不一致 ⇒ "
        "两处登记互相矛盾，列面期望值取自哪一行变得不确定"
    )
    width = _col_index(right) - _col_index(left) + 1
    assert width > 0, f"{where}: `cells`={cells!r} 跨度非正"
    return width


@lru_cache(maxsize=1)
def _contract() -> Mapping[str, Any]:
    assert _CONTRACT.is_file(), f"契约清单缺失：{_CONTRACT}（守卫必须打红而非跳过）"
    doc = json.loads(_CONTRACT.read_text(encoding="utf-8"))
    assert isinstance(doc.get("sheets"), Mapping), "清单 `sheets` 段形态非法"
    return doc


@lru_cache(maxsize=1)
def _x3_entries() -> dict[str, Mapping[str, Any]]:
    """作业面 —— 判据 = 条目登记了 ``key_family``（与实现 ``_select_x3_entries`` 同口径）。"""
    sheets = _contract()["sheets"]
    return {
        code: entry
        for code, entry in sorted(sheets.items())
        if isinstance(entry, Mapping) and entry.get("key_family")
    }


def _column_source(code: str) -> Mapping[str, Any]:
    entry = _x3_entries()[code]
    source = ((entry.get("provenance") or {}).get("column_source")) or {}
    assert isinstance(source, Mapping) and source, (
        f"{code}: 清单缺 `provenance.column_source` ⇒ 源模板取值溯源丢失，"
        "期望值无从取得（R3.7 要求 openpyxl 直读源模板）"
    )
    return source


def _snapshot(paths: Sequence[Path]) -> dict[str, tuple[int, int]]:
    snap: dict[str, tuple[int, int]] = {}
    for path in paths:
        st_ = path.stat()
        snap[path.relative_to(_TEMPLATES_DIR).as_posix()] = (st_.st_size, st_.st_mtime_ns)
    return snap


def _norm_cells(row: tuple[Any, ...], width: int) -> tuple[str, ...]:
    out = ["" if v is None else str(v) for v in row[:width]]
    out += [""] * (width - len(out))
    return tuple(out)


@lru_cache(maxsize=1)
def _live() -> _LiveBundle:
    """16 张 X-3 的源模板实读（表头行 + 示例行）+ 读取前后的只读自检快照。"""
    codes = list(_x3_entries())
    assert codes, (
        "清单里没有任何条目登记 `key_family` ⇒ 作业面为空、下面全部 ∀ 属性恒真（空转）。"
        f"真源：{_CONTRACT}"
    )

    paths: dict[str, Path] = {}
    for code in codes:
        source = _column_source(code)
        raw = source.get("file")
        assert isinstance(raw, str) and raw, f"{code}: `column_source.file` 缺失"
        path = (_REPO / raw).resolve()
        assert path.is_file(), f"{code}: 登记的源模板不存在：{path}"
        assert _TEMPLATES_DIR in path.parents, (
            f"{code}: 登记的源模板 {path} 不在 {_TEMPLATES_DIR} 下 —— "
            "R3.1 要求列头取自 backend/wp_templates/（参考副本已落后，不可作真源）"
        )
        paths[code] = path

    before = _snapshot(list(paths.values()))

    templates: dict[str, _Live] = {}
    for code, path in paths.items():
        source = _column_source(code)
        header_row = source.get("header_row")
        assert isinstance(header_row, int) and header_row >= 1, (
            f"{code}: `column_source.header_row` 形态非法: {header_row!r}"
        )
        width = _span_width(str(source.get("cells")), header_row, code)
        tab = impl.sheet_spec(code).sheet_name  # R3.8 真源 = catalog sheet_name
        wb = load_workbook(path, read_only=True, data_only=True)
        try:
            assert tab in wb.sheetnames, (
                f"{code}: catalog sheet_name {tab!r} 不在源模板 tab 列表里（实有 "
                f"{wb.sheetnames}）⇒ R3.8 的『与源模板 tab 名逐字一致』已不成立，"
                "期望值取不到"
            )
            ws = wb[tab]
            rows = {
                idx: _norm_cells(row, width)
                for idx, row in enumerate(
                    ws.iter_rows(
                        min_row=header_row, max_row=header_row + 1, values_only=True
                    ),
                    start=header_row,
                )
            }
        finally:
            wb.close()
        labels = rows.get(header_row, ())
        templates[code] = _Live(
            sheet_code=code,
            path=path,
            tab=tab,
            header_row=header_row,
            width=width,
            labels=labels,
            sample_row=rows.get(header_row + 1, tuple([""] * width)),
        )

    after = _snapshot(list(paths.values()))
    return _LiveBundle(templates=templates, before=before, after=after)


@lru_cache(maxsize=1)
def _expected_face() -> tuple[str, ...]:
    """16 张共同的列面（实读一致才有单一期望值）。"""
    variants = {live.labels for live in _live().templates.values()}
    assert len(variants) == 1, (
        "16 张 X-3 表头行不再逐字一致，不存在单一列面期望值：\n"
        + "\n".join(
            f"  {code}: {live.labels}" for code, live in sorted(_live().templates.items())
        )
    )
    face = next(iter(variants))
    assert face and all(isinstance(x, str) and x.strip() for x in face), (
        f"实读列面含空标签，逐字比对会假红：{face}"
    )
    return face


@lru_cache(maxsize=1)
def _codes() -> tuple[str, ...]:
    return tuple(_live().templates)


def _spec(code: str) -> Any:
    return impl.sheet_spec(code)


def _placeholder_positions(code: str) -> tuple[int, ...]:
    """占位位 = ``field_keys`` 的 ``None`` 位（**不假定只有 F 列**）。"""
    return tuple(i for i, key in enumerate(_spec(code).field_keys) if key is None)


# ═══════════════════════════════════════════════════════════════════════════
# 2. 判据（纯函数；期望值只有一份 = 实读列面 + catalog sheet 名）
# ═══════════════════════════════════════════════════════════════════════════

_V_SHEET_NAME = "sheet名不符"
_V_EXTRA = "多列"
_V_MISSING = "缺列"
_V_DUP = "重复列"
_V_WIDTH = "列数不等"
_V_ORDER = "列序不一致"
_V_ROW_WIDTH = "行宽不等"
_V_PLACEHOLDER = "占位列未留空"
_V_MISPLACED = "单元格错位"
_V_ROW_COUNT = "行数不等"
_V_PREFILL = "模板出现数据行"

_BLANK = object()


class _Product(NamedTuple):
    """导出产物的可判读形态：(工作表名, 逐行单元格)。``grid[0]`` 是列头行。"""

    title: str
    grid: tuple[tuple[Any, ...], ...]


def _extract(wb: Any, sheet_name: str) -> _Product:
    """从 workbook 取出产物 —— 唯一的产物读取处（in-memory 与字节往返共用）。"""
    titles = list(wb.sheetnames)
    assert titles, "导出 workbook 没有任何工作表"
    title = titles[0]
    ws = wb[title] if title == sheet_name or sheet_name not in titles else wb[sheet_name]
    grid = tuple(tuple(row) for row in ws.iter_rows(values_only=True))
    return _Product(title=title, grid=grid)


def _header_violations(product: _Product, expected: Sequence[str], sheet_name: str) -> list[str]:
    """产物列头 / sheet 名 vs 期望（实读列面 + catalog sheet 名）。"""
    problems: list[str] = []
    if product.title != sheet_name:
        problems.append(
            f"{_V_SHEET_NAME}（R3.8）：导出工作表名 {product.title!r} != catalog sheet_name "
            f"{sheet_name!r}"
        )
    if not product.grid:
        problems.append(f"{_V_WIDTH}：产物没有列头行")
        return problems

    exp = list(expected)
    cand = ["" if v is None else str(v) for v in product.grid[0]]
    extra = [c for c in cand if c not in exp]
    if extra:
        problems.append(f"{_V_EXTRA}（R3.6：源模板表头行不存在的列）：{extra}")
    missing = [e for e in exp if e not in cand]
    if missing:
        problems.append(f"{_V_MISSING}（R3.2：列集须与源模板表头行等势）：{missing}")
    dup = sorted({c for c in cand if cand.count(c) > 1})
    if dup:
        problems.append(f"{_V_DUP}：{dup}")
    if len(cand) != len(exp):
        problems.append(f"{_V_WIDTH}：期望 {len(exp)} 列、实测 {len(cand)} 列")
    if not problems and cand != exp:
        problems.append(
            f"{_V_ORDER}："
            + " / ".join(
                f"#{i + 1} 期望 {e!r} 实测 {c!r}"
                for i, (e, c) in enumerate(zip(exp, cand))
                if e != c
            )
        )
    return problems


def _expected_cell(value: Any) -> Any:
    """输入行某字段 → 期望单元格（与实现的 ``export_row_by_keys`` 语义独立复述）。"""
    if value is None:
        return _BLANK
    if isinstance(value, bool):
        return round(int(value), 2)
    if isinstance(value, (int, float)):
        return round(value, 2)
    text = str(value)
    return _BLANK if text == "" else text


def _cell_matches(actual: Any, expected: Any, *, loose: bool) -> bool:
    if expected is _BLANK:
        if actual is None:
            return True
        return isinstance(actual, str) and (actual.strip() == "" if loose else actual == "")
    if isinstance(expected, (int, float)) and not isinstance(expected, bool):
        if isinstance(actual, bool) or not isinstance(actual, (int, float)):
            return False
        return abs(float(actual) - float(expected)) <= 1e-6 * max(1.0, abs(float(expected)))
    if not isinstance(actual, str):
        return False
    return actual.strip() == str(expected).strip() if loose else actual == str(expected)


def _row_violations(
    product: _Product,
    code: str,
    expected: Sequence[str],
    input_rows: Sequence[Mapping[str, Any]],
    *,
    loose: bool,
) -> list[str]:
    """数据行：行数 / 行宽 / 占位位留空 / 非占位位对位取值。"""
    problems: list[str] = []
    field_keys = _spec(code).field_keys
    body = product.grid[1:]
    if len(body) != len(input_rows):
        problems.append(f"{_V_ROW_COUNT}：期望 {len(input_rows)} 行、实测 {len(body)} 行")
    for idx, (cells, source) in enumerate(zip(body, input_rows), start=1):
        if len(cells) != len(expected):
            problems.append(
                f"{_V_ROW_WIDTH}：第 {idx} 行 {len(cells)} 个单元格，期望 {len(expected)} 个"
                f"（占位列必须写空串保列序，不得跳过）"
            )
            continue
        for pos, (cell, key) in enumerate(zip(cells, field_keys)):
            if key is None:
                if not _cell_matches(cell, _BLANK, loose=loose):
                    problems.append(
                        f"{_V_PLACEHOLDER}：第 {idx} 行第 {pos + 1} 列"
                        f"（{expected[pos]!r}）应留空，实测 {cell!r}"
                    )
                continue
            want = _expected_cell(source.get(key))
            if not _cell_matches(cell, want, loose=loose):
                problems.append(
                    f"{_V_MISPLACED}：第 {idx} 行第 {pos + 1} 列（{expected[pos]!r} ⇐ "
                    f"{key!r}）期望 {'空' if want is _BLANK else want!r}，实测 {cell!r}"
                )
    return problems


def _template_violations(product: _Product, code: str) -> list[str]:
    """R3.1 补充：模板态**不预填**数据行（design E13：示例行预填会被当用户数据导回库）。"""
    body = product.grid[1:]
    if not body:
        return []
    sample = _live().templates[code].sample_row
    hits = [
        idx
        for idx, cells in enumerate(body, start=1)
        if any(
            isinstance(c, str) and c != "" and c in sample
            for c in cells
        )
    ]
    return [
        f"{_V_PREFILL}：模板应只含列头行，实测多出 {len(body)} 行"
        + (f"，其中第 {hits} 行含源模板示例行的取值" if hits else "")
        + f"；首个多出行={body[0]!r}"
    ]


# ═══════════════════════════════════════════════════════════════════════════
# 3. 生成器
# ═══════════════════════════════════════════════════════════════════════════

#: 单元格取值的字符面 —— 常用汉字 + 拉丁字母数字 + 中英文标点/符号。
#: 刻意**不**按列头标签拼字符（那会让人误以为期望值藏在这里）：期望值只来自实读列面。
#: 不含 ``\r`` / ``\n`` / 制表符与前后空白（xlsx 的 XML 会做空白归一，属序列化噪声）；
#: 不含 ``=``（openpyxl 对以 ``=`` 开头的字符串按公式落盘，同样是与列面无关的噪声）。
_TEXT_ALPHABET = (
    "甲乙丙丁戊己庚辛壬癸一二三四五六七八九十零壹贰叁肆伍陆柒捌玖"
    "京沪粤苏浙鲁豫川渝陕甘青蒙桂黔滇藏宁新"
    + string.ascii_letters
    + string.digits
    + "，。、；：（）【】《》“”‘’—～·%&#@!?+-*/_|^$<>"
)

#: 长中文（远超源模板列宽的极端文本）
_LONG_CN = "壹贰叁肆伍陆柒捌玖拾京沪粤苏浙鲁豫川渝陕"

_arb_text = st.text(alphabet=_TEXT_ALPHABET, min_size=0, max_size=24)
_arb_long = st.builds(lambda n: _LONG_CN * n, st.integers(min_value=3, max_value=8))
_arb_number = st.one_of(
    st.integers(min_value=-10**9, max_value=10**9),
    st.floats(
        min_value=-1e12, max_value=1e12, allow_nan=False, allow_infinity=False
    ),
    # 极小有限小数（round(…, 2) 后落 0.0 / -0.0，字节往返后可能变 int）
    st.floats(min_value=-1e-6, max_value=1e-6, allow_nan=False, allow_infinity=False),
)
_arb_cell_value = st.one_of(st.none(), _arb_text, _arb_long, _arb_number)


@lru_cache(maxsize=1)
def _all_field_keys() -> frozenset[str]:
    keys: set[str] = set()
    for code in _codes():
        keys.update(k for k in _spec(code).field_keys if k)
    return frozenset(keys)


@lru_cache(maxsize=1)
def _extra_keys() -> tuple[str, ...]:
    """「源模板无对应列」的多余字段名 —— 取清单 ``frontend_extra_fields`` 的并集减去全部列字段。

    再加两个合成名，保证即使清单某天把 extra 清空，本维度仍有实例（不空转）。
    """
    extras: set[str] = set()
    for code in _codes():
        extras.update(_spec(code).extra_fields)
        extras.add(_spec(code).entry_type_field)
    extras -= _all_field_keys()
    extras |= {"__x3prop_unknown_a", "__x3prop_unknown_b"}
    return tuple(sorted(extras))


class _AbstractRow(NamedTuple):
    """与列面同长的抽象行：逐列取值 + 缺失列集 + 多余字段。

    与具体 sheet 解耦 ⇒ 同一个生成值可materialize到 16 张各自的 ``field_keys`` 上，
    从而让「∀ sheet × ∀ 行数据」两维在同一次迭代里都覆盖到。
    """

    values: tuple[Any, ...]
    dropped: frozenset[int]
    extras: tuple[tuple[str, Any], ...]


def _arb_abstract_row(width: int) -> st.SearchStrategy[_AbstractRow]:
    return st.builds(
        _AbstractRow,
        values=st.tuples(*([_arb_cell_value] * width)),
        dropped=st.frozensets(st.integers(min_value=0, max_value=width - 1), max_size=width),
        extras=st.lists(
            st.tuples(st.sampled_from(_extra_keys()), _arb_cell_value),
            max_size=3,
        ).map(lambda pairs: tuple(dict(pairs).items())),
    )


def _arb_rows(width: int, *, max_rows: int = 5) -> st.SearchStrategy[list[_AbstractRow]]:
    return st.lists(_arb_abstract_row(width), min_size=0, max_size=max_rows)


def _materialize(code: str, rows: Sequence[_AbstractRow]) -> list[dict[str, Any]]:
    """抽象行 → 该 sheet 的输入行（占位列天然无字段；``dropped`` 位的字段整个缺失）。"""
    field_keys = _spec(code).field_keys
    out: list[dict[str, Any]] = []
    for row in rows:
        payload: dict[str, Any] = {}
        for pos, key in enumerate(field_keys):
            if key is None or pos in row.dropped:
                continue
            payload[key] = row.values[pos]
        payload.update(dict(row.extras))
        out.append(payload)
    return out


# ═══════════════════════════════════════════════════════════════════════════
# 4. 前置锚点（反空转；红 = 期望值侧或作业面出了问题，下面的属性结论不可解读）
# ═══════════════════════════════════════════════════════════════════════════



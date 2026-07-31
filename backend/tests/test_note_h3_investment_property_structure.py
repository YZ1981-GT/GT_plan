"""附注投资性房地产章节结构守卫（§五、21 上市 / §八、21 国企）。

锁定 `fix_note_h3_investment_property_structure.py` 的对齐结果：

- 表数/表名序（国企第 3 表已从重名 `以公允价值计量` 改为「未办妥产权证书的投资性房地产」）
- 国企两级表头（以成本计量 7 列 / 以公允价值计量 8 列；期初·期末 rowspan=2 无 group）
- 上市列转置 flat 5 列
- columns key 与前端 `h3NoteSectionMap.ts` 逐字一致（防 seed/push 漂移）
- 行集层次（国企 5 层 / 3 层，每层 1 合计 + 2 类别）
- 无 `header_label` 假行、无「可无限量添加行」占位行、无重名
- openpyxl 直读源 xlsx tab 名与两级子表头交叉比对
- 反向自检（防守卫空转）

spec: h3-investment-property-disclosure-alignment (Task 4)
"""
from __future__ import annotations

import importlib.util
import json
from pathlib import Path

import openpyxl
import pytest

_ROOT = Path(__file__).resolve().parent.parent
_FIX = _ROOT / "scripts" / "fix" / "fix_note_h3_investment_property_structure.py"
_SRC_XLSX = _ROOT / "wp_templates" / "H" / "H3 投资性房地产.xlsx"
_FE_MAP = (
    _ROOT.parent / "audit-platform" / "frontend" / "src" / "components"
    / "workpaper" / "composables" / "h3NoteSectionMap.ts"
)


def _load_fix():
    spec = importlib.util.spec_from_file_location("_fix_h3_struct", _FIX)
    mod = importlib.util.module_from_spec(spec)
    spec.loader.exec_module(mod)
    return mod


FIX = _load_fix()

_LISTED_TRANSPOSED_KEYS = ["label", "房屋、建筑物", "土地使用权", "在建工程", "合计"]

_EXPECTED_KEYS = {
    "listed": {
        "按成本计量的投资性房地产": _LISTED_TRANSPOSED_KEYS,
        "按公允价值计量的投资性房地产": _LISTED_TRANSPOSED_KEYS,
        "未办妥产权证书的情况": ["label", "book_value", "reason"],
    },
    "soe": {
        "以成本计量": [
            "label", "begin", "buy_or_provision", "transfer_in",
            "disposal", "transfer_out", "end",
        ],
        "以公允价值计量": [
            "label", "begin_fair", "buy", "transfer_in", "fair_change_pl",
            "disposal", "transfer_out", "end_fair",
        ],
        "未办妥产权证书的投资性房地产": ["label", "book_value", "reason"],
    },
}

# 两级表（分组数 2）与其无 group 的 rowspan=2 列
_TWO_LEVEL = {
    "listed": {},
    "soe": {"以成本计量": {"begin", "end"}, "以公允价值计量": {"begin_fair", "end_fair"}},
}


def _section(variant: str) -> dict:
    path, section_number, _plan, _texts = FIX._TARGETS[variant]
    doc = json.loads(Path(path).read_text(encoding="utf-8"))
    return next(s for s in doc["sections"] if str(s.get("section_number")) == section_number)


@pytest.mark.parametrize("variant", ["listed", "soe"])
def test_check_passes(variant: str):
    _changes, warnings, errs = FIX._runner(variant, dry_run=True, check=True)
    assert not errs, f"{variant} 结构欠账：{errs}"
    assert not warnings, f"{variant} 告警：{warnings}"


@pytest.mark.parametrize("variant", ["listed", "soe"])
def test_table_names_unique_and_ordered(variant: str):
    names = [str(t.get("name")) for t in _section(variant).get("tables") or []]
    assert names == FIX.EXPECTED[variant], f"{variant} 表序/表名漂移：{names}"
    assert len(names) == len(set(names)), f"{variant} 表名重复（同名会互相覆盖丢整表）：{names}"


def test_soe_third_table_renamed():
    """国企第 3 表不得再与第 2 表同名 `以公允价值计量`。"""
    names = [str(t.get("name")) for t in _section("soe").get("tables") or []]
    assert names[2] == "未办妥产权证书的投资性房地产", names
    assert names.count("以公允价值计量") == 1, names


@pytest.mark.parametrize("variant", ["listed", "soe"])
def test_columns_keys_and_groups(variant: str):
    by_name = {str(t.get("name")): t for t in _section(variant).get("tables") or []}
    for name, want_keys in _EXPECTED_KEYS[variant].items():
        tbl = by_name[name]
        cols = tbl.get("columns") or []
        assert [c.get("key") for c in cols] == want_keys, f"{variant}/{name} 列 key 漂移"
        assert cols[0].get("is_label"), f"{variant}/{name} 首列未 is_label"
        assert str(cols[0].get("label")) == str(tbl["headers"][0]), (
            f"{variant}/{name} columns[0].label ≠ headers[0]"
        )
        assert str(tbl.get("guidance") or "").strip(), f"{variant}/{name} 缺 guidance"
        no_group = _TWO_LEVEL[variant].get(name)
        if no_group:
            groups = tbl.get("_column_groups") or []
            assert len(groups) == 2, f"{variant}/{name} 应有 2 分组，实得 {groups}"
            for c in cols[1:]:
                if c["key"] in no_group:
                    assert not c.get("group"), f"{variant}/{name} {c['key']} 是 rowspan=2 列不应带 group"
                else:
                    assert c.get("group"), f"{variant}/{name} {c['key']} 缺 group"
        else:
            assert not tbl.get("_column_groups"), f"{variant}/{name} 单级表残留 _column_groups"
            assert any(c.get("flat") for c in cols), f"{variant}/{name} 单级表未标 flat"


@pytest.mark.parametrize("variant", ["listed", "soe"])
def test_no_fake_or_placeholder_rows(variant: str):
    for tbl in _section(variant).get("tables") or []:
        for r in tbl.get("rows") or []:
            assert r.get("row_type") != "header_label", f"{variant}/{tbl.get('name')} 残留 header_label"
            assert "可无限量添加行" not in str(r.get("label") or ""), (
                f"{variant}/{tbl.get('name')} 残留「可无限量添加行」占位行"
            )


def test_soe_layer_row_structure():
    """国企 5 层 / 3 层，每层 1 合计行 + 2 类别行。"""
    by_name = {str(t.get("name")): t for t in _section("soe").get("tables") or []}
    cost_rows = by_name["以成本计量"]["rows"]
    assert len(cost_rows) == 15, len(cost_rows)
    assert [r["label"] for r in cost_rows if r.get("is_total")] == [
        "一、账面原值合计", "二、累计折旧和累计摊销合计", "三、投资性房地产账面净值合计",
        "四、投资性房地产减值准备累计金额合计", "五、投资性房地产账面价值合计",
    ]
    fair_rows = by_name["以公允价值计量"]["rows"]
    assert len(fair_rows) == 9, len(fair_rows)
    assert [r["label"] for r in fair_rows if r.get("is_total")] == [
        "一、成本合计", "二、公允价值变动合计", "三、投资性房地产账面价值合计",
    ]
    # 每层下辖两类别
    assert sum(1 for r in cost_rows if "房屋、建筑物" in str(r["label"])) == 5
    assert sum(1 for r in cost_rows if "土地使用权" in str(r["label"])) == 5


def test_source_xlsx_two_level_headers():
    """openpyxl 直读源 xlsx：国企两级子表头与模板叶子列名一致。"""
    wb = openpyxl.load_workbook(_SRC_XLSX, data_only=True)
    assert "附注披露信息（上市公司）" in wb.sheetnames
    soe = wb["附注披露信息（国有企业）"]
    # 以成本计量：行 9 子表头（「购置或」+ 行 10「计提」合成「购置或计提」）
    assert [soe.cell(9, c).value for c in range(3, 7)] == [
        "购置或", "自用房地产或存货转入", "处 置", "转为自用房地产",
    ]
    assert soe.cell(10, 3).value == "计提"
    # 以公允价值计量：行 30 子表头
    assert [soe.cell(30, c).value for c in range(3, 8)] == [
        "购置", "自用房地产或存货转入", "公允价值变动损益", "处 置", "转为自用房地产",
    ]
    # 上市列转置表头（行 8）
    listed = wb["附注披露信息（上市公司）"]
    assert [listed.cell(8, c).value for c in range(2, 6)] == [
        "房屋、建筑物", "土地使用权", "在建工程", "合计",
    ]


def test_frontend_map_section_numbers():
    """前端映射章节号必须是 五、21 / 八、21，且不得出现 八、22（那是固定资产）。"""
    src = _FE_MAP.read_text(encoding="utf-8")
    assert "listed: '五、21'" in src, "前端 listed 章节号漂移"
    assert "soe: '八、21'" in src, "前端 soe 章节号漂移"
    assert "'八、22'" not in src, "八、22 是固定资产章节，H3 不得引用"


def test_reverse_self_check():
    """validate_section 能抓出缺 columns / 重名 / header_label（防守卫空转）。"""
    from _note_structure_kit import validate_section

    broken = {
        "tables": [
            {
                "name": "以成本计量",
                "headers": ["项目", "期初余额"],
                "rows": [{"label": "项目", "row_type": "header_label"}],
                "guidance": "x",
            },
            {"name": "以成本计量", "headers": ["项目"], "rows": [], "guidance": "y"},
        ],
    }
    errs = validate_section(broken, ["以成本计量"])
    assert any("缺 columns" in e for e in errs), errs
    assert any("header_label" in e for e in errs), errs
    assert any("表名重复" in e for e in errs), errs

"""Unit tests for prefilled-download include_guidance=true feature.

Tests that guidance annotations are injected as blue italic text alongside
placeholders in the exported docx template.

**Validates: Requirements 12.2, 12.3, 12.4, 12.6**
"""

from __future__ import annotations

import json
import os
import shutil
import tempfile
from pathlib import Path
from unittest.mock import patch

import pytest
from docx import Document
from docx.shared import Pt, RGBColor

from app.services.wp_docx_template_parser import parse_template


# ─── Helpers ──────────────────────────────────────────────────────────────────

BLUE = RGBColor(0x00, 0x70, 0xC0)


def _create_template_with_placeholders() -> str:
    """Create a temp docx with known placeholders. Returns path."""
    doc = Document()
    doc.add_heading("审计业务约定书", level=1)
    doc.add_paragraph("甲方：${entity_name:被审计单位}")
    doc.add_paragraph("审计期间：${audit_period:审计期间}")
    doc.add_paragraph("报告日期：${report_date:报告日期}")

    table = doc.add_table(rows=2, cols=2)
    table.cell(0, 0).text = "项目"
    table.cell(0, 1).text = "内容"
    table.cell(1, 0).text = "委托单位"
    table.cell(1, 1).text = "${client_org:委托单位}"

    fd, path = tempfile.mkstemp(suffix=".docx")
    os.close(fd)
    doc.save(path)
    return path


def _create_guidance_json(wp_code: str, tmp_dir: str) -> Path:
    """Create a temporary guidance JSON file."""
    guidance = {
        "wp_code": wp_code,
        "title": "测试模板编制指引",
        "sections": [
            {
                "title": "本表用途",
                "content": "这是测试用的编制指引说明。",
            },
            {
                "heading": "编制步骤",
                "content": "第一步：填写甲方信息\n第二步：确认审计期间",
            },
        ],
        "recommended_questions": ["测试问题1？"],
    }
    guidance_dir = Path(tmp_dir) / "wp_guidance"
    guidance_dir.mkdir(parents=True, exist_ok=True)
    guidance_path = guidance_dir / f"{wp_code}.json"
    with open(guidance_path, "w", encoding="utf-8") as f:
        json.dump(guidance, f, ensure_ascii=False)
    return guidance_dir.parent


def _inject_guidance(template_path: str, guidance_data_dir: Path | None = None) -> str:
    """Apply the guidance injection logic (mirrors endpoint logic).

    Returns path to the modified docx.
    """
    from docx.shared import Pt, RGBColor
    from app.services.wp_docx_template_parser import parse_template as _parse_tpl

    wp_code = "A8-1"

    # Copy template
    fd, tmp_path = tempfile.mkstemp(suffix=".docx")
    os.close(fd)
    shutil.copy2(template_path, tmp_path)

    # Load guidance JSON
    guidance_data: dict = {}
    if guidance_data_dir:
        guidance_path = guidance_data_dir / "wp_guidance" / f"{wp_code}.json"
        if guidance_path.exists():
            with open(guidance_path, "r", encoding="utf-8") as gf:
                guidance_data = json.load(gf)

    # Parse template
    structure = _parse_tpl(tmp_path)

    # Build annotations
    guidance_annotations: dict[str, str] = {}
    for placeholder in structure.placeholders:
        annotation = f"【说明: {placeholder.label}"
        if placeholder.data_type == "date":
            annotation += "，格式: YYYY年MM月DD日"
        elif placeholder.data_type == "number":
            annotation += "，填写数字"
        annotation += "】"
        guidance_annotations[placeholder.pattern] = annotation

    # Inject annotations
    doc = Document(tmp_path)
    guidance_injected = False

    def _inject_in_paragraph(para):
        nonlocal guidance_injected
        for pattern, annotation_text in guidance_annotations.items():
            if pattern in para.text:
                run = para.add_run(f" {annotation_text}")
                run.font.italic = True
                run.font.color.rgb = RGBColor(0x00, 0x70, 0xC0)
                run.font.size = Pt(9)
                guidance_injected = True

    for para in doc.paragraphs:
        _inject_in_paragraph(para)

    for table in doc.tables:
        for row in table.rows:
            for cell in row.cells:
                for para in cell.paragraphs:
                    _inject_in_paragraph(para)

    # Append sections if guidance exists
    if guidance_data and guidance_data.get("sections"):
        doc.add_paragraph()
        heading_run = doc.add_paragraph().add_run("── 模板编制指引 ──")
        heading_run.font.italic = True
        heading_run.font.color.rgb = RGBColor(0x00, 0x70, 0xC0)
        heading_run.font.size = Pt(10)
        for section in guidance_data["sections"]:
            sec_title = section.get("title") or section.get("heading") or ""
            sec_content = section.get("content") or ""
            if sec_title:
                title_para = doc.add_paragraph()
                title_run = title_para.add_run(sec_title)
                title_run.font.bold = True
                title_run.font.italic = True
                title_run.font.color.rgb = RGBColor(0x00, 0x70, 0xC0)
                title_run.font.size = Pt(9)
            if sec_content:
                content_para = doc.add_paragraph()
                content_run = content_para.add_run(sec_content)
                content_run.font.italic = True
                content_run.font.color.rgb = RGBColor(0x00, 0x70, 0xC0)
                content_run.font.size = Pt(9)
        guidance_injected = True

    if guidance_injected:
        doc.save(tmp_path)

    return tmp_path


# ─── Tests ────────────────────────────────────────────────────────────────────


class TestIncludeGuidanceAnnotations:
    """Tests that include_guidance injects blue italic annotations."""

    def test_guidance_annotations_injected_next_to_placeholders(self):
        """Placeholder paragraphs get blue italic annotation runs appended."""
        template_path = _create_template_with_placeholders()
        try:
            result_path = _inject_guidance(template_path)
            doc = Document(result_path)

            # Check paragraph with ${entity_name:被审计单位}
            para_text = doc.paragraphs[1].text  # "甲方：${entity_name:被审计单位} 【说明: 被审计单位】"
            assert "【说明: 被审计单位】" in para_text

            # Check the annotation run styling
            runs = doc.paragraphs[1].runs
            # Last run should be the annotation
            annotation_run = runs[-1]
            assert annotation_run.font.italic is True
            assert annotation_run.font.color.rgb == BLUE
            assert annotation_run.font.size == Pt(9)

            os.unlink(result_path)
        finally:
            os.unlink(template_path)

    def test_guidance_date_field_includes_format_hint(self):
        """Date-type placeholders get format guidance in their annotation."""
        template_path = _create_template_with_placeholders()
        try:
            result_path = _inject_guidance(template_path)
            doc = Document(result_path)

            # report_date is detected as date type
            full_text = "\n".join(p.text for p in doc.paragraphs)
            assert "格式: YYYY年MM月DD日" in full_text

            os.unlink(result_path)
        finally:
            os.unlink(template_path)

    def test_guidance_injected_in_table_cells(self):
        """Placeholders in table cells also get guidance annotations."""
        template_path = _create_template_with_placeholders()
        try:
            result_path = _inject_guidance(template_path)
            doc = Document(result_path)

            # Check table cell with ${client_org:委托单位}
            table_text = ""
            for table in doc.tables:
                for row in table.rows:
                    for cell in row.cells:
                        table_text += cell.text + "\n"

            assert "【说明: 委托单位】" in table_text

            os.unlink(result_path)
        finally:
            os.unlink(template_path)

    def test_guidance_blue_italic_style_preserved(self):
        """All annotation runs have consistent blue italic 9pt styling."""
        template_path = _create_template_with_placeholders()
        try:
            result_path = _inject_guidance(template_path)
            doc = Document(result_path)

            annotation_runs = []
            for para in doc.paragraphs:
                for run in para.runs:
                    if "【说明:" in run.text:
                        annotation_runs.append(run)

            assert len(annotation_runs) >= 3  # entity_name, audit_period, report_date

            for run in annotation_runs:
                assert run.font.italic is True, f"Run '{run.text}' not italic"
                assert run.font.color.rgb == BLUE, f"Run '{run.text}' not blue"
                assert run.font.size == Pt(9), f"Run '{run.text}' not 9pt"

            os.unlink(result_path)
        finally:
            os.unlink(template_path)

    def test_original_placeholders_preserved(self):
        """Guidance does NOT replace placeholders — they stay in the document."""
        template_path = _create_template_with_placeholders()
        try:
            result_path = _inject_guidance(template_path)
            doc = Document(result_path)

            full_text = "\n".join(p.text for p in doc.paragraphs)
            assert "${entity_name:被审计单位}" in full_text
            assert "${audit_period:审计期间}" in full_text
            assert "${report_date:报告日期}" in full_text

            os.unlink(result_path)
        finally:
            os.unlink(template_path)

    def test_guidance_sections_appended_when_json_exists(self):
        """When guidance JSON exists, sections are appended at document end."""
        template_path = _create_template_with_placeholders()
        tmp_dir = tempfile.mkdtemp()
        try:
            data_dir = _create_guidance_json("A8-1", tmp_dir)
            result_path = _inject_guidance(template_path, guidance_data_dir=data_dir)
            doc = Document(result_path)

            full_text = "\n".join(p.text for p in doc.paragraphs)
            assert "── 模板编制指引 ──" in full_text
            assert "本表用途" in full_text
            assert "这是测试用的编制指引说明。" in full_text
            assert "编制步骤" in full_text

            os.unlink(result_path)
        finally:
            os.unlink(template_path)
            shutil.rmtree(tmp_dir)

    def test_guidance_sections_styling_blue_italic(self):
        """Appended guidance sections also use blue italic styling."""
        template_path = _create_template_with_placeholders()
        tmp_dir = tempfile.mkdtemp()
        try:
            data_dir = _create_guidance_json("A8-1", tmp_dir)
            result_path = _inject_guidance(template_path, guidance_data_dir=data_dir)
            doc = Document(result_path)

            # Find the "── 模板编制指引 ──" heading
            guidance_heading_found = False
            for para in doc.paragraphs:
                for run in para.runs:
                    if "模板编制指引" in run.text:
                        assert run.font.italic is True
                        assert run.font.color.rgb == BLUE
                        guidance_heading_found = True
                        break

            assert guidance_heading_found, "Guidance heading not found in document"

            os.unlink(result_path)
        finally:
            os.unlink(template_path)
            shutil.rmtree(tmp_dir)

    def test_no_guidance_json_still_injects_field_annotations(self):
        """When no guidance JSON exists, field-level annotations still work."""
        template_path = _create_template_with_placeholders()
        try:
            # No guidance_data_dir → no JSON loaded
            result_path = _inject_guidance(template_path, guidance_data_dir=None)
            doc = Document(result_path)

            full_text = "\n".join(p.text for p in doc.paragraphs)
            # Field annotations should still be present
            assert "【说明: 被审计单位】" in full_text
            # But no guidance sections at the end
            assert "── 模板编制指引 ──" not in full_text

            os.unlink(result_path)
        finally:
            os.unlink(template_path)

    def test_non_placeholder_paragraphs_unchanged(self):
        """Paragraphs without placeholders are not modified."""
        template_path = _create_template_with_placeholders()
        try:
            result_path = _inject_guidance(template_path)
            doc = Document(result_path)

            # First paragraph is the heading "审计业务约定书" — should be unchanged
            heading = doc.paragraphs[0]
            assert heading.text == "审计业务约定书"
            # Should not have guidance annotations
            assert "【说明:" not in heading.text

            os.unlink(result_path)
        finally:
            os.unlink(template_path)

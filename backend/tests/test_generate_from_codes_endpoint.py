"""按编码生成底稿端点守卫。

spec formula-governance-coverage-closure R5 走查撞上的缺陷（2026-10-22）：原端点在函数体内 import 全仓从未实现的
``populate_parsed_data`` ⇒ 对**任何**调用都在入口 ImportError 500；前端「生成底稿」在模板集生成 5xx 时回退到本端点，
回退同样失败。既有 property 测试只在内存里模拟端点逻辑、从未真发请求，所以一直全绿。

spec formula-fetch-and-wp-generation-defect-closure R5（2026-10-23）：模板改与链路编排同一解析
（``init_workpaper_from_template``：模板目录索引 + 多文件合并 + 场景裁剪）。原先按 ``gt_template_library.json``
精确匹配，模板集 1025 个编码里 722 个不在库里（E1 / D2 等多页底稿全在其中）⇒ 只生成叫「底稿E1」的三行空白册。

本守卫按真实路由发请求（SQLite 真 ORM，底稿存储目录指到临时目录），断言生成文件的真实页签。
"""
from __future__ import annotations

import uuid
from decimal import Decimal

import openpyxl
import pytest
import sqlalchemy as sa
from fastapi import FastAPI
from httpx import ASGITransport, AsyncClient
from sqlalchemy.dialects.sqlite.base import SQLiteTypeCompiler
from sqlalchemy.ext.asyncio import AsyncSession, async_sessionmaker, create_async_engine
from sqlalchemy.pool import StaticPool

SQLiteTypeCompiler.visit_JSONB = SQLiteTypeCompiler.visit_JSON

import app.models  # noqa: E402, F401  全量注册（外键引用的表要在 metadata 里）
from app.core.database import get_db  # noqa: E402
from app.deps import get_current_user  # noqa: E402
from app.models.audit_platform_models import AccountCategory, TrialBalance  # noqa: E402
from app.models.base import Base, PermissionLevel, ProjectStatus, ProjectUserRole, UserRole  # noqa: E402
from app.models.core import Project, ProjectUser, User  # noqa: E402
from app.models.procedure_models import ProcedureInstance  # noqa: E402
from app.models.wp_generation_models import WorkpaperGenerationItem, WorkpaperGenerationOperation  # noqa: E402
from app.models.workpaper_models import WorkingPaper, WpIndex  # noqa: E402
from app.routers import wp_template  # noqa: E402
from app.services import wp_parsed_data_service, wp_template_init_service  # noqa: E402
from app.services.wp_generate_file import BLANK_NOTICE, BLANK_NOTICE_ROW  # noqa: E402
from app.services.wp_parsed_data_service import ensure_parsed_data_placeholder  # noqa: E402

pytestmark = pytest.mark.asyncio

#: 生成端点要求本项目 edit 级成员（project_users）并读取粗裁结果（procedure_instances，读不到即 503）
_TABLES = [User.__table__, Project.__table__, ProjectUser.__table__, ProcedureInstance.__table__,
           TrialBalance.__table__, WorkpaperGenerationOperation.__table__, WorkpaperGenerationItem.__table__,
           WpIndex.__table__, WorkingPaper.__table__]
_URL = "/api/projects/{pid}/working-papers/generate-from-codes"

#: E1 主文件（E1-1至E1-11）与另两个常规程序文件（E1-14至15 / E1-18至23）各取一个页签；IPO 文件（E1-26至32）只在
#: 非 normal 场景合并；「(修订前)」历史页签永不合并。取自 backend/wp_templates/E 的四个 E1 文件。
_E1_MAIN_SHEET = "货币资金审定表E1-1"
_E1_OTHER_SHEETS = ("货币资金分析表E1-14", "银行存款截止测试表E1-21")
_E1_IPO_SHEET = "现金交易分析E1-26"
_E1_HISTORICAL_SHEET = "货币资金分析表F1-6 (修订前)"


@pytest.fixture
async def world(tmp_path, monkeypatch):
    monkeypatch.chdir(tmp_path)  # 旧实现按相对路径 storage/projects/... 落文件：断言新实现不再往那里写
    storage = tmp_path / "wp_storage"
    monkeypatch.setattr(wp_template_init_service, "STORAGE_DIR", storage)
    engine = create_async_engine("sqlite+aiosqlite:///:memory:", poolclass=StaticPool)
    async with engine.begin() as conn:
        await conn.run_sync(lambda c: Base.metadata.create_all(c, tables=_TABLES))
    async with async_sessionmaker(engine, class_=AsyncSession, expire_on_commit=False)() as db:
        user = User(id=uuid.uuid4(), username="gen_user", email="gen_user@example.com", hashed_password="x",
                    role=UserRole.auditor, is_active=True)
        with_tb, without_tb, ipo = uuid.uuid4(), uuid.uuid4(), uuid.uuid4()
        db.add(user)
        for pid, name, scenario in ((with_tb, "【测试】按编码生成_有试算表", "normal"),
                                    (without_tb, "【测试】按编码生成_无试算表", "normal"),
                                    (ipo, "【测试】按编码生成_IPO", "ipo")):
            db.add(Project(id=pid, name=name, client_name=name, status=ProjectStatus.execution,
                           template_type="soe", report_scope="standalone", audit_year=2025, scenario=scenario))
        await db.flush()
        for pid in (with_tb, without_tb, ipo):
            db.add(ProjectUser(project_id=pid, user_id=user.id, role=ProjectUserRole.auditor,
                               permission_level=PermissionLevel.edit))
        for pid in (with_tb, ipo):
            db.add(TrialBalance(project_id=pid, year=2025, company_code="001", standard_account_code="1001",
                                account_name="库存现金", account_category=AccountCategory.asset,
                                unadjusted_amount=Decimal("100"), audited_amount=Decimal("100")))
        await db.commit()

        api = FastAPI()
        api.include_router(wp_template.router)

        async def _db():
            yield db

        async def _user() -> User:
            return user

        api.dependency_overrides[get_db] = _db
        api.dependency_overrides[get_current_user] = _user
        async with AsyncClient(transport=ASGITransport(app=api), base_url="http://gen") as http:
            yield {"http": http, "db": db, "with_tb": with_tb, "without_tb": without_tb, "ipo": ipo,
                   "root": tmp_path, "storage": storage}
    await engine.dispose()


async def _papers(db: AsyncSession, pid: uuid.UUID) -> list[tuple[str, str, str, dict | None, uuid.UUID]]:
    rows = (await db.execute(
        sa.select(WpIndex.wp_code, WpIndex.wp_name, WorkingPaper.file_path, WorkingPaper.parsed_data,
                  WorkingPaper.id)
        .join(WorkingPaper, WorkingPaper.wp_index_id == WpIndex.id)
        .where(WpIndex.project_id == pid)
        .order_by(WpIndex.wp_code)
    )).all()
    return [tuple(r) for r in rows]


def _sheetnames(path: str) -> list[str]:
    wb = openpyxl.load_workbook(path, read_only=True)
    try:
        return list(wb.sheetnames)
    finally:
        wb.close()


async def _generate(world, pid_key: str, codes: list[str]):
    r = await world["http"].post(_URL.format(pid=world[pid_key]), json={"wp_codes": codes, "year": 2025})
    assert r.status_code == 200, r.text
    return r.json()


async def test_e1_gets_real_multi_file_template_and_source_name(world):
    """E1：真实模板（三个常规程序文件合并、不带 IPO 文件与历史页签），名称取平台名称源，文件进底稿存储目录。"""
    body = await _generate(world, "with_tb", ["E1"])
    assert body["created"] == 1 and body["failures"] == [] and body["created_codes"] == ["E1"]
    assert body["blank_codes"] == []
    [(code, name, file_path, parsed, wp_id)] = await _papers(world["db"], world["with_tb"])
    assert (code, name) == ("E1", "货币资金"), "名称应取 wp_account_mapping（不再叫「底稿E1」）"
    expected = world["storage"] / str(world["with_tb"]) / f"{wp_id}.xlsx"
    assert file_path == str(expected) and expected.exists()
    sheets = _sheetnames(file_path)
    assert _E1_MAIN_SHEET in sheets and set(_E1_OTHER_SHEETS) <= set(sheets), sheets
    assert _E1_IPO_SHEET not in sheets, "normal 场景不应合并 IPO / 上市 / 重组 / 舞弊应对文件"
    assert _E1_HISTORICAL_SHEET not in sheets
    assert parsed["_created_by"] == "generate_from_codes" and parsed["wp_code"] == "E1"
    assert not (world["root"] / "storage").exists(), "不应再往旧的 storage/projects 相对路径写空白册"


async def test_ipo_project_merges_ipo_file(world):
    """项目场景 ipo ⇒ 与链路编排同口径，合并 IPO / 上市 / 重组 / 舞弊应对文件。"""
    await _generate(world, "ipo", ["E1"])
    [(_c, _n, file_path, _p, _id)] = await _papers(world["db"], world["ipo"])
    assert _E1_IPO_SHEET in _sheetnames(file_path)


async def test_code_without_template_is_reported_blank(world):
    """找不到模板 ⇒ 仍建空白底稿，但单列 blank_codes、消息写明，空白册里有说明文字。"""
    body = await _generate(world, "with_tb", ["ZZ9", "E1"])
    assert body["created"] == 2 and body["blank_codes"] == ["ZZ9"]
    assert "没有找到底稿模板" in body["message"] and "ZZ9" in body["message"]
    papers = {p[0]: p for p in await _papers(world["db"], world["with_tb"])}
    wb = openpyxl.load_workbook(papers["ZZ9"][2])
    try:
        assert wb.sheetnames == ["ZZ9"]
        assert wb.active[f"A{BLANK_NOTICE_ROW}"].value == BLANK_NOTICE
    finally:
        wb.close()
    assert papers["ZZ9"][1] == "底稿ZZ9", "没有名称源时才回退「底稿{编码}」"


async def test_word_template_kept_as_docx(world):
    """主模板是 Word（B60）⇒ 原样复制为 .docx，不当成工作簿去写表头。"""
    body = await _generate(world, "with_tb", ["B60"])
    assert body["blank_codes"] == [] and body["failures"] == []
    [(_c, _n, file_path, _p, wp_id)] = await _papers(world["db"], world["with_tb"])
    assert file_path.endswith(f"{wp_id}.docx")


async def test_failure_after_file_written_rolls_back_and_discards_file(world, monkeypatch):
    """落文件之后失败 ⇒ 该底稿行回滚、文件删掉，记入 failures；同批其他编码照常生成。"""
    original = wp_parsed_data_service.ensure_parsed_data_placeholder

    def _boom(wp, *, wp_code, created_by):
        if wp_code == "H1":
            raise RuntimeError("占位写入失败")
        return original(wp, wp_code=wp_code, created_by=created_by)

    monkeypatch.setattr(wp_parsed_data_service, "ensure_parsed_data_placeholder", _boom)
    body = await _generate(world, "with_tb", ["H1", "E1"])
    assert body["created_codes"] == ["E1"]
    assert [f["wp_code"] for f in body["failures"]] == ["H1"]
    papers = await _papers(world["db"], world["with_tb"])
    assert [p[0] for p in papers] == ["E1"]
    files = sorted(p.name for p in (world["storage"] / str(world["with_tb"])).iterdir())
    assert files == [f"{papers[0][4]}.xlsx"], "失败的 H1 不应留下孤儿文件"


async def test_unknown_request_field_is_rejected(world):
    """请求体有未登记字段 ⇒ 422（不再静默丢弃，前后端字段漂移当场暴露）。"""
    r = await world["http"].post(_URL.format(pid=world["with_tb"]),
                                 json={"wp_codes": ["E1"], "year": 2025, "selected_templates": ["E1"]})
    assert r.status_code == 422
    assert await _papers(world["db"], world["with_tb"]) == []


async def test_generate_twice_skips_existing(world):
    url = _URL.format(pid=world["with_tb"])
    assert (await world["http"].post(url, json={"wp_codes": ["E1"], "year": 2025})).status_code == 200
    again = (await world["http"].post(url, json={"wp_codes": ["E1"], "year": 2025})).json()
    assert again["created"] == 0 and again["skipped"] == 1 and again["skipped_codes"] == ["E1"]
    assert len(await _papers(world["db"], world["with_tb"])) == 1


async def test_operation_id_keeps_per_item_semantics_and_replays(world):
    operation_id = str(uuid.uuid4())
    payload = {"wp_codes": ["ZZ9", "E1"], "year": 2025, "operation_id": operation_id}
    first = await world["http"].post(_URL.format(pid=world["with_tb"]), json=payload)
    second = await world["http"].post(_URL.format(pid=world["with_tb"]), json=payload)
    assert first.status_code == second.status_code == 200
    assert first.json() == second.json()
    assert first.json()["created_codes"] == ["ZZ9", "E1"]
    assert first.json()["blank_codes"] == ["ZZ9"]
    assert first.json()["skipped"] == 0 and first.json()["operation_id"] == operation_id
    assert len(await _papers(world["db"], world["with_tb"])) == 2


async def test_generate_without_trial_balance_is_rejected(world):
    r = await world["http"].post(_URL.format(pid=world["without_tb"]), json={"wp_codes": ["E1"], "year": 2025})
    assert r.status_code == 422
    assert "试算表" in str(r.json())
    assert await _papers(world["db"], world["without_tb"]) == []


async def test_placeholder_keeps_existing_parsed_data():
    wp = WorkingPaper(project_id=uuid.uuid4(), wp_index_id=uuid.uuid4(), file_path="x.xlsx")
    wp.parsed_data = {"html_data": {"Sheet1": {"cells": {"A1": 1}}}}
    assert ensure_parsed_data_placeholder(wp, wp_code="E1", created_by="generate_from_codes") is False
    assert wp.parsed_data == {"html_data": {"Sheet1": {"cells": {"A1": 1}}}}
    blank = WorkingPaper(project_id=uuid.uuid4(), wp_index_id=uuid.uuid4(), file_path="y.xlsx")
    assert ensure_parsed_data_placeholder(blank, wp_code="D2", created_by="standard_conversion") is True
    assert blank.parsed_data["_created_by"] == "standard_conversion" and blank.parsed_data["html_data"] == {}


# ── 表头（2026-10-26：合并格模板表头静默填不上） ─────────────────────────────────────────


def _cell(path: str, sheet: str, coord: str):
    wb = openpyxl.load_workbook(path)
    try:
        return wb[sheet][coord].value
    finally:
        wb.close()


async def test_merged_cell_templates_get_header_and_no_header_failures(world):
    """J2A（表头页「底稿目录」）与 A10（活动页标签、值区是合并格）原先写表头时 AttributeError，回执照报成功。"""
    body = await _generate(world, "with_tb", ["J2A", "A10"])
    assert sorted(body["created_codes"]) == ["A10", "J2A"] and body["failures"] == []
    assert body["header_failures"] == [] and "表头未填上" not in body["message"]
    papers = {p[0]: p for p in await _papers(world["db"], world["with_tb"])}
    client = "【测试】按编码生成_有试算表"
    assert client in str(_cell(papers["J2A"][2], "底稿目录", "A2"))
    assert client in str(_cell(papers["A10"][2], "审计程序A10", "A3"))


async def test_header_fill_exception_is_reported_in_receipt(world, monkeypatch):
    """填表头抛错 ⇒ 底稿照常生成，但回执 header_failures 与 message 写明（原先只记 warning）。"""
    from app.services import wp_header_service

    async def _boom(**_kwargs):
        raise RuntimeError("模拟：表头写入崩溃")

    monkeypatch.setattr(wp_header_service, "fill_workpaper_header", _boom)
    body = await _generate(world, "with_tb", ["E1"])
    assert body["created_codes"] == ["E1"] and body["failures"] == []
    assert body["header_failures"] == [{"wp_code": "E1", "reason": "填写表头时出错，详见服务日志"}]
    assert "1 个底稿表头未填上：E1（填写表头时出错，详见服务日志）" in body["message"]
    assert "模拟" not in body["message"], "异常原文只进日志"


async def test_external_link_template_header_is_filled(world):
    """D2-5 的表头格原是外部链接（=[4]底稿目录!$A$2）：procedure-chain-followup-closure 5.3 时填不进去，只能在回执里
    写明原因；spec post-push-residual-defect-closure R3 起按缓存值归类改成本册引用 / 静态标签 ⇒ 表头填上本项目值，
    回执不再报（仍有填不上的外部标签引用时照报，见 test_wp_header_fill.test_unresolvable_external_header_is_still_reported）。"""
    body = await _generate(world, "with_tb", ["D2-5"])
    assert body["created_codes"] == ["D2-5"] and body["failures"] == []
    assert body["header_failures"] == [] and "表头未填上" not in body["message"]
    [(_c, _n, file_path, _p, _id)] = await _papers(world["db"], world["with_tb"])
    wb = openpyxl.load_workbook(file_path, data_only=True)  # 按缓存值读（与 HTML 渲染同口径）
    try:
        ws = wb.active
        header = [str(c.value) for row in ws.iter_rows(min_row=1, max_row=10, max_col=12) for c in row
                  if c.value is not None]
    finally:
        wb.close()
    assert any("【测试】按编码生成_有试算表" in v for v in header), header


async def test_template_without_header_is_not_resaved(world):
    """A30（电子底稿归档检查表）没有可填的表头 ⇒ 文件与模板逐字节一致（不重存，openpyxl 重存会丢数据验证扩展）。"""
    import hashlib

    from app.services.wp_template_finder import find_all_template_files, find_template_file_any

    template = find_template_file_any("A30")
    assert template is not None and find_all_template_files("A30") == [template], "前提：A30 是单文件模板（不合并）"
    body = await _generate(world, "with_tb", ["A30"])
    assert body["header_failures"] == []
    [(_c, _n, file_path, _p, _id)] = await _papers(world["db"], world["with_tb"])
    digest = lambda p: hashlib.sha256(open(p, "rb").read()).hexdigest()  # noqa: E731
    assert digest(file_path) == digest(template)

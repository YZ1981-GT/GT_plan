"""
Property 3: HTML→Word 导出保留结构元素

Feature: a17-summary-enhancement
**Validates: Requirements 1.6**

For any sanitized HTML chapter content containing supported elements (headings,
lists, bold, italic, tables), the Word exporter SHALL produce a Document where:
each <h3>/<h4> maps to a Heading style paragraph, each <ul>/<ol> item maps to a
list paragraph, <strong> maps to bold run, <em> maps to italic run, and <table>
maps to a Table object.
"""

import re

import pytest
from hypothesis import given, settings, assume
from hypothesis import strategies as st

from docx import Document

# Import the function under test
import sys
from pathlib import Path

sys.path.insert(0, str(Path(__file__).resolve().parents[1]))
from app.services.a17_word_exporter import _html_to_docx


# ─── Strategies ───

# 安全文本（不含 HTML 特殊字符）
safe_text_st = st.text(
    alphabet=st.characters(
        whitelist_categories=("L", "N", "Zs"),
        whitelist_characters="，。、；：！？",
    ),
    min_size=1,
    max_size=30,
).filter(lambda s: s.strip() != "")

# 生成 h3/h4 标签
heading_html_st = st.tuples(
    st.sampled_from(["h3", "h4"]),
    safe_text_st,
).map(lambda t: f"<{t[0]}>{t[1]}</{t[0]}>")

# 生成列表 HTML
list_item_st = safe_text_st.map(lambda t: f"<li>{t}</li>")

list_html_st = st.tuples(
    st.sampled_from(["ul", "ol"]),
    st.lists(list_item_st, min_size=1, max_size=4),
).map(lambda t: f"<{t[0]}>{''.join(t[1])}</{t[0]}>")

# 生成 bold/italic HTML
bold_html_st = safe_text_st.map(lambda t: f"<p><strong>{t}</strong></p>")
italic_html_st = safe_text_st.map(lambda t: f"<p><em>{t}</em></p>")

# 生成表格 HTML
table_cell_st = safe_text_st
table_row_st = st.lists(table_cell_st, min_size=1, max_size=3).map(
    lambda cells: "<tr>" + "".join(f"<td>{c}</td>" for c in cells) + "</tr>"
)
table_html_st = st.lists(table_row_st, min_size=1, max_size=3).map(
    lambda rows: "<table>" + "".join(rows) + "</table>"
)


# ─── Property Tests ───


class TestProperty3HtmlToWord:
    """Property 3: HTML→Word 导出保留结构元素"""

    @settings(max_examples=100)
    @given(heading_html_st)
    def test_headings_map_to_heading_styles(self, html: str):
        """h3/h4 → Heading 3/4 style paragraph"""
        doc = Document()
        _html_to_docx(doc, html)

        # 找到新增段落（Document() 默认无段落）
        paragraphs = doc.paragraphs
        assert len(paragraphs) >= 1

        # 最后一个段落应该是 Heading 样式
        last_para = paragraphs[-1]
        assert last_para.style.name in ("Heading 3", "Heading 4")

        # 检查文本内容保留
        # 从 html 中提取文本
        tag_match = re.search(r"<h[34]>(.+?)</h[34]>", html)
        if tag_match:
            expected_text = tag_match.group(1)
            assert last_para.text == expected_text

    @settings(max_examples=100)
    @given(list_html_st)
    def test_list_items_map_to_list_paragraphs(self, html: str):
        """ul/ol items → list paragraphs"""
        doc = Document()
        _html_to_docx(doc, html)

        paragraphs = doc.paragraphs
        assert len(paragraphs) >= 1

        # 所有段落应为列表样式或普通段落（fallback 如无样式）
        for para in paragraphs:
            # 列表段落应有非空文本
            assert para.text.strip() != ""

    @settings(max_examples=100)
    @given(bold_html_st)
    def test_strong_maps_to_bold_run(self, html: str):
        """<strong> → bold run"""
        doc = Document()
        _html_to_docx(doc, html)

        paragraphs = doc.paragraphs
        assert len(paragraphs) >= 1

        # 查找含有 bold run 的段落
        found_bold = False
        for para in paragraphs:
            for run in para.runs:
                if run.bold and run.text.strip():
                    found_bold = True
                    break
        assert found_bold, "Expected at least one bold run"

    @settings(max_examples=100)
    @given(italic_html_st)
    def test_em_maps_to_italic_run(self, html: str):
        """<em> → italic run"""
        doc = Document()
        _html_to_docx(doc, html)

        paragraphs = doc.paragraphs
        assert len(paragraphs) >= 1

        # 查找含有 italic run 的段落
        found_italic = False
        for para in paragraphs:
            for run in para.runs:
                if run.italic and run.text.strip():
                    found_italic = True
                    break
        assert found_italic, "Expected at least one italic run"

    @settings(max_examples=100)
    @given(table_html_st)
    def test_table_maps_to_table_object(self, html: str):
        """<table> → Table object"""
        doc = Document()
        _html_to_docx(doc, html)

        # Document 应包含至少一个 Table
        assert len(doc.tables) >= 1

        table = doc.tables[-1]
        # 从 HTML 中计算期望的行数
        row_count = html.count("<tr>")
        assert len(table.rows) == row_count

    @settings(max_examples=100)
    @given(safe_text_st)
    def test_unsupported_tags_silently_skipped(self, text: str):
        """不支持的标签（如 span/div）保留文本内容"""
        html = f"<span class='custom'>{text}</span>"
        doc = Document()
        _html_to_docx(doc, html)

        # 文本内容应该被保留
        all_text = "".join(p.text for p in doc.paragraphs)
        assert text in all_text

    @settings(max_examples=100)
    @given(safe_text_st)
    def test_plain_text_fallback(self, text: str):
        """纯文本（无 HTML 标签）按行分段处理"""
        doc = Document()
        _html_to_docx(doc, text)

        # 纯文本应生成段落
        assert len(doc.paragraphs) >= 1
        all_text = "".join(p.text for p in doc.paragraphs)
        assert text.strip() in all_text

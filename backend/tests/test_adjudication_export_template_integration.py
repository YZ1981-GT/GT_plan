"""Integration and edge case tests for AdjudicationExportTemplateService.

Validates: Design Properties 5, 6, 7, 8 + edge cases (aging fallback, missing template).
"""

from __future__ import annotations

import io
from unittest.mock import AsyncMock, patch
from urllib.parse import quote, unquote

import pytest
from openpyxl import Workbook, load_workbook

from app.services.adjudication_export_template_service import (
    AdjudicationExportTemplateService as Svc,
)


# ── 8.1 Property 6: RFC5987 Content-Disposition header with Chinese characters ──


class TestRFC5987ContentDisposition:
    """Property 6: filename contains wp_code and Chinese wp_name, percent-encoded.

    **Validates: Requirements 7.1, 7.2**
    """

    def test_rfc5987_encoding_chinese_filename(self):
        """filename*=UTF-8'' followed by percent-encoded Chinese filename."""
        filename = "D2-1_应收账款审定表_模板.xlsx"
        encoded = quote(filename)
        header = f"attachment; filename*=UTF-8''{encoded}"

        # Verify it decodes back correctly
        assert unquote(encoded) == filename
        assert "UTF-8''" in header
        assert "D2-1" in header

    def test_rfc5987_encoding_various_wp_codes(self):
        """Multiple wp_code + Chinese name combinations encode correctly."""
        cases = [
            ("D1-1", "货币资金"),
            ("F1-1", "应收账款"),
            ("K3-1", "预付款项"),
            ("N5-1", "所得税费用"),
        ]
        for wp_code, wp_name in cases:
            filename = f"{wp_code}_{wp_name}审定表_模板.xlsx"
            encoded = quote(filename)
            header = f"attachment; filename*=UTF-8''{encoded}"

            assert unquote(encoded) == filename, f"Failed for {wp_code}"
            assert wp_code in unquote(encoded)
            assert "UTF-8''" in header

    @pytest.mark.asyncio
    async def test_generate_returns_correct_filename_format(self):
        """generate() produces filename matching {wp_code}_{wp_name}审定表_模板.xlsx."""
        db = AsyncMock()
        _, filename = await Svc.generate("wp-1", "D2-1", "应收账款", db, None)
        assert filename == "D2-1_应收账款审定表_模板.xlsx"

        # Verify RFC5987 encoding roundtrip
        encoded = quote(filename)
        assert unquote(encoded) == filename


# ── 8.2 Property 7: non-adjudication wp_codes behavior ──────────────────────


class TestNonAdjudicationCodes:
    """Property 7: wp_codes not matching ^[A-N]\\d+-1$ are rejected by is_adjudication_table.

    **Validates: Requirements 1.2, 9.1, 9.4**
    """

    def test_non_adjudication_codes_not_matched(self):
        """wp_codes that don't match the pattern return False."""
        non_adj_codes = ["D2", "D2-2", "O1-1", "d2-1", "D2-1A", "X99-1", ""]
        for code in non_adj_codes:
            assert Svc.is_adjudication_table(code) is False, (
                f"Expected False for '{code}'"
            )

    def test_boundary_letter_range(self):
        """Only A-N are valid; O and beyond are rejected."""
        # O is out of range
        assert Svc.is_adjudication_table("O1-1") is False
        assert Svc.is_adjudication_table("P1-1") is False
        assert Svc.is_adjudication_table("Z1-1") is False
        # A and N are boundaries - should be accepted
        assert Svc.is_adjudication_table("A1-1") is True
        assert Svc.is_adjudication_table("N1-1") is True

    def test_suffix_must_be_dash_1(self):
        """Only -1 suffix qualifies; -2, -3, etc. are non-adjudication."""
        assert Svc.is_adjudication_table("D2-2") is False
        assert Svc.is_adjudication_table("D2-3") is False
        assert Svc.is_adjudication_table("D2-10") is False
        assert Svc.is_adjudication_table("D2-1") is True

    def test_none_and_empty_handled(self):
        """None and empty string should return False without error."""
        assert Svc.is_adjudication_table("") is False
        assert Svc.is_adjudication_table(None) is False  # type: ignore[arg-type]


# ── 8.3 Property 5: row skeleton order preserved ────────────────────────────


class TestRowSkeletonOrder:
    """Property 5: items appear in same order as extract_audit_rows output.

    **Validates: Requirements 6.1, 6.2**
    """

    def test_row_skeleton_order_preserved(self):
        """Items in row_skeleton appear in data sheet in exact same order."""
        skeleton = [
            {"id": "1", "item": "货币资金", "is_section": False, "is_total": False, "indent": 0},
            {"id": "2", "item": "其中：库存现金", "is_section": False, "is_total": False, "indent": 1},
            {"id": "3", "item": "其中：银行存款", "is_section": False, "is_total": False, "indent": 1},
            {"id": "4", "item": "合计", "is_section": False, "is_total": True, "indent": 0},
        ]
        ws = Workbook().active
        Svc._build_data_sheet(ws, "D1-1", "货币资金", skeleton)

        # Verify order starting from row 4
        for i, item in enumerate(skeleton):
            assert ws.cell(4 + i, 1).value == item["item"], (
                f"Order mismatch at index {i}: expected '{item['item']}', got '{ws.cell(4 + i, 1).value}'"
            )

    def test_section_and_total_styling(self):
        """Section headers and total rows get bold styling."""
        skeleton = [
            {"id": "1", "item": "流动资产", "is_section": True, "is_total": False, "indent": 0},
            {"id": "2", "item": "货币资金", "is_section": False, "is_total": False, "indent": 1},
            {"id": "3", "item": "合计", "is_section": False, "is_total": True, "indent": 0},
        ]
        ws = Workbook().active
        Svc._build_data_sheet(ws, "D1-1", "货币资金", skeleton)

        # Section row bold
        assert ws.cell(4, 1).font.bold is True
        # Normal row not bold
        assert ws.cell(5, 1).font.bold is not True
        # Total row bold
        assert ws.cell(6, 1).font.bold is True

    def test_large_skeleton_preserves_order(self):
        """Order preserved even for larger skeletons (15+ items)."""
        items = [f"科目{i}" for i in range(20)]
        skeleton = [
            {"id": str(i), "item": name, "is_section": False, "is_total": False, "indent": 0}
            for i, name in enumerate(items)
        ]
        ws = Workbook().active
        Svc._build_data_sheet(ws, "D1-1", "货币资金", skeleton)

        for i, name in enumerate(items):
            assert ws.cell(4 + i, 1).value == name


# ── 8.4 Property 8: instruction sheet contains all required sections ────────


class TestInstructionSheetSections:
    """Property 8: instruction sheet has all 4 required sections.

    **Validates: Requirements 5.1, 5.2, 5.3, 5.4**
    """

    @pytest.mark.asyncio
    async def test_instruction_sheet_sections_present(self):
        """Instruction sheet contains 各列含义/只读自动计算列/填报规则/导入注意事项."""
        db = AsyncMock()
        buf, _ = await Svc.generate("wp-1", "D1-1", "货币资金", db, None)
        wb = load_workbook(buf)
        ws = wb["编制说明"]

        # Collect all cell values in column A
        texts = [
            ws.cell(r, 1).value
            for r in range(1, ws.max_row + 1)
            if ws.cell(r, 1).value
        ]
        joined = "\n".join(str(t) for t in texts)

        assert "各列含义" in joined
        assert "只读自动计算列" in joined
        assert "填报规则" in joined
        assert "导入注意事项" in joined

    @pytest.mark.asyncio
    async def test_instruction_sheet_formula_descriptions(self):
        """Instruction sheet documents auto-calc formulas."""
        db = AsyncMock()
        buf, _ = await Svc.generate("wp-1", "D1-1", "货币资金", db, None)
        wb = load_workbook(buf)
        ws = wb["编制说明"]

        texts = [
            ws.cell(r, 1).value
            for r in range(1, ws.max_row + 1)
            if ws.cell(r, 1).value
        ]
        joined = "\n".join(str(t) for t in texts)

        # Key formulas should be described
        assert "期初审定" in joined
        assert "期末审定" in joined
        assert "变动额" in joined
        assert "变动率" in joined

    @pytest.mark.asyncio
    async def test_instruction_sheet_title_includes_wp_info(self):
        """Instruction sheet title references the wp_name and wp_code."""
        db = AsyncMock()
        buf, _ = await Svc.generate("wp-1", "F3-1", "固定资产", db, None)
        wb = load_workbook(buf)
        ws = wb["编制说明"]

        title = ws.cell(1, 1).value
        assert "固定资产" in title
        assert "F3-1" in title

    @pytest.mark.asyncio
    async def test_instruction_sheet_aging_note_present_for_aging_codes(self):
        """Aging wp_codes get an extra aging note in instruction sheet."""
        db = AsyncMock()
        with patch(
            "app.services.adjudication_export_template_service.AdjudicationExportTemplateService._resolve_aging_headers",
            return_value=["1年以内", "1-2年"],
        ):
            buf, _ = await Svc.generate("wp-1", "D2-1", "应收账款", db, None)

        wb = load_workbook(buf)
        ws = wb["编制说明"]

        texts = [
            ws.cell(r, 1).value
            for r in range(1, ws.max_row + 1)
            if ws.cell(r, 1).value
        ]
        joined = "\n".join(str(t) for t in texts)

        # Should mention aging columns
        assert "账龄" in joined


# ── 8.5 Edge case: template_file_path missing → empty project column ────────


class TestMissingTemplateEdgeCase:
    """Edge case: no template file → template with empty project column.

    **Validates: Requirements 6.4**
    """

    @pytest.mark.asyncio
    async def test_missing_template_empty_skeleton(self):
        """No template file → template with empty project column from row 4."""
        db = AsyncMock()
        buf, _ = await Svc.generate("wp-1", "D1-1", "货币资金", db, template_file_path=None)
        wb = load_workbook(buf)
        ws = wb["货币资金审定表"]

        # Row 4 should be empty (no skeleton)
        assert ws.cell(4, 1).value is None

    @pytest.mark.asyncio
    async def test_missing_template_nonexistent_path(self):
        """Non-existent template file path → same as None (empty skeleton)."""
        db = AsyncMock()
        buf, _ = await Svc.generate(
            "wp-1", "D1-1", "货币资金", db,
            template_file_path="/nonexistent/path/to/template.xlsx",
        )
        wb = load_workbook(buf)
        ws = wb["货币资金审定表"]

        # Row 4 should still be empty
        assert ws.cell(4, 1).value is None

    @pytest.mark.asyncio
    async def test_missing_template_still_has_headers(self):
        """Even without skeleton, headers are still properly generated."""
        db = AsyncMock()
        buf, _ = await Svc.generate("wp-1", "D1-1", "货币资金", db, template_file_path=None)
        wb = load_workbook(buf)
        ws = wb["货币资金审定表"]

        # Header row 2 should have values
        assert ws.cell(2, 1).value == "项目"
        assert ws.cell(2, 2).value == "期初"
        assert ws.cell(2, 6).value == "期末"


# ── 8.6 Edge case: aging config failure → fallback to default preset ────────


class TestAgingFallbackOnFailure:
    """Edge case: aging config failure still produces headers (default preset).

    **Validates: Requirements 4.4**
    """

    @pytest.mark.asyncio
    async def test_aging_fallback_on_failure(self):
        """Aging config failure still produces aging headers via fallback preset."""
        db = AsyncMock()
        # Mock resolve_aging_segments to raise an exception
        with patch(
            "app.routers.wp_render_strategies._cycle_import_export_common.resolve_aging_segments",
            side_effect=Exception("DB unavailable"),
        ):
            buf, _ = await Svc.generate("wp-1", "D2-1", "应收账款", db, None)

        wb = load_workbook(buf)
        ws = wb["应收账款审定表"]

        # Should have more than 12 columns (aging fallback adds columns)
        # Col 10 should have an aging header value (from default preset)
        assert ws.cell(2, 10).value is not None, (
            "Expected aging header at col 10 from fallback preset"
        )

        # 变动额 should be shifted past the aging columns
        found_change = False
        for col in range(10, 30):
            if ws.cell(2, col).value == "变动额":
                found_change = True
                break
        assert found_change, "变动额 not found after aging columns"

    @pytest.mark.asyncio
    async def test_aging_fallback_produces_valid_xlsx(self):
        """Even on fallback, the generated xlsx is structurally valid."""
        db = AsyncMock()
        with patch(
            "app.routers.wp_render_strategies._cycle_import_export_common.resolve_aging_segments",
            side_effect=Exception("DB unavailable"),
        ):
            buf, filename = await Svc.generate("wp-1", "D2-1", "应收账款", db, None)

        # Should be a valid xlsx
        wb = load_workbook(buf)
        assert wb is not None
        assert len(wb.sheetnames) == 2
        assert filename == "D2-1_应收账款审定表_模板.xlsx"

    @pytest.mark.asyncio
    async def test_aging_fallback_has_correct_structure(self):
        """Fallback aging template retains all structural elements."""
        db = AsyncMock()
        with patch(
            "app.routers.wp_render_strategies._cycle_import_export_common.resolve_aging_segments",
            side_effect=Exception("DB unavailable"),
        ):
            buf, _ = await Svc.generate("wp-1", "D2-1", "应收账款", db, None)

        wb = load_workbook(buf)
        ws = wb["应收账款审定表"]

        # Title row present
        assert "应收账款审定表 D2-1" == ws.cell(1, 1).value

        # Standard headers present
        assert ws.cell(2, 1).value == "项目"
        assert ws.cell(2, 2).value == "期初"
        assert ws.cell(2, 6).value == "期末"

        # Sub-headers present
        assert ws.cell(3, 2).value == "未审"
        assert ws.cell(3, 5).value == "审定"

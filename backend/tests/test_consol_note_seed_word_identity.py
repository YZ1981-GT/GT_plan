"""合并附注种子与 Word 模板逐表身份一致（spec consol-note-header-library-and-node-reaggregate R7.6 / design §16.6）。

- 国企：第五章 ``五-N-k`` = 国企合并 Word ``##SECTION:八、N##`` 节内第 k 张表（物理顺序，含没有表标记的表）。原种子造过
  假章节 五-10（父章节「续：」），其后每节错后一节；Word 同步工具又按位置把错表的表名 / 表头写进去。逐表钉住列、表头路径、
  行名、父章节、表名（有表标记的取 Word 表前段落）、表后文字（同步工具 Phase 6 同一口径）。
- 上市：Word 缺 五、71~78 节标记（同步工具 ``KNOWN_SECTION_GAPS`` 登记），这 20 张表都落在 五、70 节标记下；种子编号与
  单体模板一致不改号，真表位置按表名逐张核对过（``LISTED_TAIL``）。被标成 五、71/73/74 的表就是种子 五-79/81/82。
- 同步工具（长期可重跑）：身份对不上的不写；节标记后的文字属于下一节；重跑对两版种子都没有可写的。
"""
from __future__ import annotations

import json
import re
from collections import Counter, defaultdict
from functools import lru_cache
from pathlib import Path

import pytest
from docx import Document

from app.services.consol_note_formula_service import multi_header_to_column_groups
from app.services.note_table_header_library import column_paths, norm
from app.services.note_word_table import _row_keys, heading_key, parse_word_table, rows_disjoint, seed_table_structure
from scripts.seed import sync_note_templates_from_word as S

DATA = Path(__file__).resolve().parents[1] / "data"
DOCX = DATA / "audit_report_templates" / "disclosure_notes"
SID = re.compile(r"^五-(\d+)-(\d+)$")

#: 上市合并 Word 五、70 节标记下的 20 张表（物理顺序）对应的种子章节
LISTED_TAIL = ("五-70-1", "五-71-1", "五-71-2", "五-72-1", "五-73-1", "五-74-1", "五-74-2", "五-75-1", "五-76-1",
               "五-77-1", "五-77-2", *(f"五-78-{k}" for k in range(1, 10)))
#: 同步工具错加的重复表（与 五-79-3~7、五-82-3~6 是同一张 Word 表），R7.6 已删
LISTED_DROPPED = (*(f"五-71-{k}" for k in range(3, 8)), *(f"五-74-{k}" for k in range(3, 7)))
#: 同步工具重跑时身份对不上、待修的条目（R7.8：上市 五-10-2~8 内容错位，修好后删本条）
KNOWN_MISMATCHES = {"soe": [], "listed": ["[listed] 五-10-6 ← Word 五、10:5：行名没有交集"]}


@lru_cache(maxsize=None)
def _word(std: str) -> tuple[dict, ...]:
    return tuple(S.extract_all_word_tables(DOCX / f"{std}_consolidated.docx"))


def _seed(std: str) -> list[dict]:
    return json.loads((DATA / f"consol_note_sections_{std}.json").read_bytes().decode("utf-8"))


def _positions(std: str, chapter: str) -> dict[str, dict]:
    """Word ``##SECTION:{chapter}、N##`` 节内每张表的物理位置 → 种子章节号 五-N-k。"""
    out: dict[str, dict] = {}
    ordinal: Counter = Counter()
    for wt in _word(std):
        head, _, num = (wt.get("section") or "").partition("、")
        if head == chapter and num.isdigit():
            ordinal[int(num)] += 1
            out[f"五-{int(num)}-{ordinal[int(num)]}"] = wt
    return out


def _structure_diffs(entry: dict, wt: dict) -> list[str]:
    """列数、逐列表头路径（Word 标签列表头留白时不比标签列）、行名。"""
    word = parse_word_table(wt["el"], "")
    if len(entry.get("headers") or []) != word.columns:
        return [f"列数 {len(entry.get('headers') or [])} ≠ Word {word.columns}"]
    out = []
    mine, theirs = column_paths(entry), [tuple(norm(x) for x in p) for p in word.header_paths]
    if mine[1:] != theirs[1:] or (theirs[0] and mine[0] != theirs[0]):
        out.append("表头路径")
    if _row_keys([r[0] if isinstance(r, list) and r else "" for r in entry.get("rows") or []]) != _row_keys(word.row_labels):
        out.append("行名")
    return out


def _sibling_titles(seed: list[dict]) -> dict[str, set[str]]:
    """同 parent_seq 的兄弟表名（不含自己）：同步工具 Phase 6 去重口径。"""
    groups: dict = defaultdict(set)
    for e in seed:
        if e.get("title"):
            groups[e.get("parent_seq", 0)].add(e["title"].strip())
    return {e["section_id"]: groups[e.get("parent_seq", 0)] - {(e.get("title") or "").strip()} for e in seed}


def test_soe_chapter_five_is_word_chapter_eight():
    """国企第五章章节集合 = Word 第八章 221 张表的位置（90 节有表；一张 Word 表恰一条种子）。"""
    pos = _positions("soe", "八")
    assert (len(pos), len({sid.rsplit("-", 1)[0] for sid in pos})) == (221, 90)
    assert {e["section_id"] for e in _seed("soe") if SID.match(e["section_id"])} == set(pos)


def test_soe_seed_tables_are_the_word_tables():
    seed = _seed("soe")
    by, siblings = {e["section_id"]: e for e in seed}, _sibling_titles(seed)
    bad = []
    for sid, wt in _positions("soe", "八").items():
        e, (n, k) = by[sid], map(int, SID.match(sid).groups())
        problems = _structure_diffs(e, wt)
        if heading_key(e.get("parent_section")) != heading_key(wt.get("section_heading")) or e.get("parent_seq") != n:
            problems.append(f"父章节「{e.get('parent_section')}」/{e.get('parent_seq')} ≠ Word 节「{wt.get('section_heading')}」")
        tag, word_title = S.parse_tag(wt["tag"]), (wt.get("title_text") or "").strip()
        if tag is not None and ((tag["section"], tag["table_idx"]) != (n, k - 1)
                                or (word_title and (e.get("title") or "").strip() != word_title)):
            problems.append(f"表标记 {wt['tag']} / 表名「{e.get('title')}」≠「{wt.get('title_text')}」")
        raw = wt.get("text_after_raw") or ""
        want = S._dedup_sibling_titles(S.clean_text_after(raw), siblings[sid]) if raw else ""
        if (e.get("text_after_table") or "") != want:
            problems.append("表后文字")
        if e.get("continuation_of") not in (None, f"五-{n}-{k - 1}"):
            problems.append(f"续表标记 {e['continuation_of']}")
        if problems:
            bad.append((sid, problems))
    assert bad == []


def test_listed_tail_follows_the_tables_under_word_section_70():
    under = [wt for wt in _word("listed") if wt.get("section") == "五、70"]
    by = {e["section_id"]: e for e in _seed("listed")}
    assert len(under) == len(LISTED_TAIL)
    assert [(sid, d) for sid, wt in zip(LISTED_TAIL, under) if (d := _structure_diffs(by[sid], wt))] == []
    assert not set(LISTED_DROPPED) & set(by)
    # 同步工具按标记错写过的 5 张：表名回到同步工具引入前的原名；表后文字取真表，到下一节节标题为止（这段没有节标记）
    assert by["五-71-1"]["title"] == "公允价值变动收益（注：以下不存在的项目可以删除）"
    assert by["五-73-1"]["title"] == "资产减值损失（损失以“—”号填列）（注：以下不存在的项目可以删除）"
    assert by["五-74-2"]["title"] == "试运行销售损益" and by["五-74-2"]["text_after_table"] == "（说明成本的具体核算方法。）"
    assert "text_after_table" not in by["五-71-2"] and "text_after_table" not in by["五-73-1"]
    tail = by["五-74-1"]["text_after_table"]
    assert tail.startswith("【提示：") and tail.endswith("】") and "试运行销售损益" not in tail


def test_listed_tables_tagged_into_missing_sections_are_seed_79_81_82():
    """Word 缺 五、71~78 节标记，其后 五、79/81/82 的表被标成 五、71/73/74：就是种子 五-79/81/82，结构逐表一致。"""
    by = {e["section_id"]: e for e in _seed("listed")}
    pairs = [(f"五-{t['section'] + 8}-{t['table_idx'] + 1}", wt) for wt in _word("listed")
             if (t := S.parse_tag(wt["tag"])) and t["chapter"] == "五" and t["section"] in (71, 73, 74)]
    assert len(pairs) == 14
    assert [(sid, _structure_diffs(by[sid], wt)) for sid, wt in pairs if _structure_diffs(by[sid], wt)] == []
    assert {(s, n) for s, n in S.KNOWN_SECTION_GAPS} == {("listed", 71), ("listed", 73), ("listed", 74)}


@pytest.mark.parametrize("std", ["soe", "listed"])
def test_word_sync_tool_has_nothing_to_write(std, monkeypatch):
    """同步工具合并模板各 Phase 重跑无可写（``--check`` 的 dry-run 只数不改，这里直接数）；身份对不上的只剩登记项。"""
    monkeypatch.setattr(S, "IDENTITY_MISMATCHES", [])
    monkeypatch.setattr(S, "KNOWN_GAP_HITS", set())
    word, data = list(_word(std)), _seed(std)
    assert S.phase1_consol(word, data, std, True, False) == (0, 0)
    assert S.phase2_consol(word, data, std, True, False) == 0
    assert S.phase3_consol(word, data, std, True, False) == (0, 0)
    assert S.phase6_text_after_and_parent(word, data, std, "consol", True, False) == (0, 0)
    assert S.check_title_drift(word, data, std, "consol", False) == []
    assert S.IDENTITY_MISMATCHES == KNOWN_MISMATCHES[std]
    assert S.KNOWN_GAP_HITS == {k for k in S.KNOWN_SECTION_GAPS if k[0] == std}  # 登记项仍对不上（未失效）


def test_text_after_stops_at_section_marker(tmp_path):
    """表后文字收集到下一个节标记为止：节标记后的节标题 / 小标题属于下一节。"""
    doc = Document()
    doc.add_paragraph("##SECTION:八、1##")
    doc.add_paragraph("货币资金")
    doc.add_paragraph("{{table:八、1:0}}货币资金")
    doc.add_table(rows=2, cols=2).cell(0, 0).text = "项目"
    doc.add_paragraph("（提示：本表注释）")
    doc.add_paragraph("##SECTION:八、2##")
    doc.add_paragraph("交易性金融资产")
    doc.add_paragraph("交易性金融资产分类")
    doc.add_table(rows=2, cols=2).cell(0, 0).text = "项目"
    path = tmp_path / "t.docx"
    doc.save(str(path))
    first, second = S.extract_all_word_tables(path)
    assert first["text_after_raw"] == "（提示：本表注释）"
    assert (second["section"], second["section_heading"], second["title_text"]) == ("八、2", "交易性金融资产", "交易性金融资产分类")


def test_identity_helpers(monkeypatch):
    monkeypatch.setattr(S, "IDENTITY_MISMATCHES", [])
    assert heading_key("（1）一年内到期的长期借款") == heading_key("一年内到期的长期借款") == "一年内到期的长期借款"
    assert heading_key("信用减值损失（损失以“—”号填列）") == "信用减值损失" == heading_key("8.信用减值损失")
    assert rows_disjoint([["库存现金"], ["合  计"]], [["应收账款"], ["合计"]]) is True
    assert rows_disjoint([["库存现金"], ["银行存款"]], [["银行存款", "1"]]) is False
    assert rows_disjoint([["合  计"]], [["应收账款"]]) is False  # 只有合计行：不判
    entry = {"parent_section": "货币资金", "rows": [["库存现金"], ["合  计"]]}
    wt = {"tag": "八、1:0", "section_heading": "货币资金", "data_rows": [["库存现金", ""]]}
    assert S.identity_mismatch(entry, wt, "soe", "五-1-1") is False
    assert S.identity_mismatch(entry, {**wt, "section_heading": "交易性金融资产"}, "soe", "五-1-1") is True
    assert S.identity_mismatch(entry, {**wt, "data_rows": [["应收账款", ""]]}, "soe", "五-1-1") is True
    assert S.identity_mismatch(entry, {**wt, "data_rows": [["应收账款", ""]]}, "soe", "五-1-1", rows=False) is False
    assert [m.split("：", 1)[1] for m in S.IDENTITY_MISMATCHES] == ["Word 节「交易性金融资产」≠ 种子父章节「货币资金」", "行名没有交集"]


def test_seed_table_structure_reproduces_correct_seed_tables():
    """平台表结构构建（同步工具新增 / 重建表共用）与位置一直正确的国企 五-1~9 逐字段相同（R7.6 前就在对的位置）。"""
    by = {e["section_id"]: e for e in _seed("soe")}
    cases = [(sid, wt) for sid, wt in _positions("soe", "八").items() if int(SID.match(sid).group(1)) <= 9]
    assert len(cases) == 47
    keys = ("headers", "multi_header", "_column_groups", "rows")
    bad = [sid for sid, wt in cases if {k: by[sid].get(k) for k in keys} != {k: seed_table_structure(wt["el"]).get(k) for k in keys}]
    assert bad == []
    mh = [by[sid] for sid, _wt in cases if by[sid].get("multi_header")]
    assert mh and all(e["_column_groups"] == multi_header_to_column_groups(e["multi_header"]) for e in mh)


def test_realign_record_covers_every_moved_section():
    """新旧章节号对照（已保存数据迁移用）：旧号各对一个新号、新号都在种子里；上市只记并掉的 9 张重复表。"""
    rec = json.loads((DATA / "consol_note_section_realign_r7.json").read_bytes().decode("utf-8"))["sources"]
    for std in ("soe", "listed"):
        sids = {e["section_id"] for e in _seed(std)}
        olds = [o for v in rec[std].values() for o in v]
        assert set(rec[std]) <= sids and len(olds) == len(set(olds)), std
    assert {o for v in rec["listed"].values() for o in v[1:]} == set(LISTED_DROPPED)
    assert rec["soe"]["五-9-12"] == ["五-9-12", "五-10-5"]  # 假章节「续：」的表并回 五-9（留 Phase 2 条目）
    assert rec["soe"]["五-40-3"][0] == "五-41-3"  # J1 引用的设定提存计划列示

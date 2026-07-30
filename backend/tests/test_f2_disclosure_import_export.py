"""F2 附注披露多区块导入导出契约测试。

Spec: .kiro/specs/f2-inventory-disclosure-template-alignment/ Task 13.4

三类断言：
1. **前后端常量镜像**（防双真源漂移）：后端 `_LISTED_CATEGORIES` / `_SOE_CATEGORIES` /
   `_DR_ROW_DEFS` / `_*_SOURCE_KEYS` 必须与前端 .ts 源码逐条一致
2. **workbook 结构**：模板 sheet 齐备、表头正确、只读区块带提示
3. **导出→导入往返**：写进去的值能原样读回来（含"空=不覆盖"语义）
"""
from __future__ import annotations

import io
import json
import re
from pathlib import Path

import pytest
from openpyxl import load_workbook

from app.routers.wp_render_strategies import _f2_disclosure_import_export as mod

FRONTEND = (
    Path(__file__).resolve().parents[2]
    / "audit-platform" / "frontend" / "src" / "components" / "workpaper" / "composables"
)


def _read(name: str) -> str:
    return (FRONTEND / name).read_text(encoding="utf-8")


def _parse_ts_categories(src: str, const_name: str) -> list[tuple[str, str]]:
    """从 .ts 源码提取 `{ rowKey: 'x', label: 'y', ... }` 序列。"""
    start = src.index(f"export const {const_name}")
    end = src.index("\n]", start)
    block = src[start:end]
    return [
        (m.group(1), m.group(2))
        for m in re.finditer(r"rowKey:\s*'([^']+)',\s*\n?\s*label:\s*'([^']+)'", block)
    ]


def _parse_ts_source_keys(src: str, const_name: str) -> dict[str, tuple[str, ...]]:
    start = src.index(f"export const {const_name}")
    end = src.index("\n]", start)
    block = src[start:end]
    out: dict[str, tuple[str, ...]] = {}
    for m in re.finditer(
        r"rowKey:\s*'([^']+)'[\s\S]*?sourceKeys:\s*\[([^\]]*)\]", block
    ):
        keys = tuple(k.strip().strip("'\"") for k in m.group(2).split(",") if k.strip())
        out[m.group(1)] = keys
    return out


# ════════════════════════════════════════════════════════════════════
# 1. 前后端常量镜像
# ════════════════════════════════════════════════════════════════════


def test_listed_categories_mirror_frontend():
    fe = _parse_ts_categories(
        _read("useF2DisclosureListed.ts"), "F2_LISTED_DISCLOSURE_CATEGORIES"
    )
    assert fe, "前端上市分类常量解析失败（格式变了？）"
    assert list(mod._LISTED_CATEGORIES) == fe


def test_soe_categories_mirror_frontend():
    fe = _parse_ts_categories(_read("useF2DisclosureSoe.ts"), "F2_SOE_DISCLOSURE_CATEGORIES")
    assert fe, "前端国企分类常量解析失败（格式变了？）"
    assert list(mod._SOE_CATEGORIES) == fe


def test_source_keys_mirror_frontend():
    fe_listed = _parse_ts_source_keys(
        _read("useF2DisclosureListed.ts"), "F2_LISTED_DISCLOSURE_CATEGORIES"
    )
    fe_soe = _parse_ts_source_keys(
        _read("useF2DisclosureSoe.ts"), "F2_SOE_DISCLOSURE_CATEGORIES"
    )
    assert fe_listed and fe_soe
    assert mod._LISTED_SOURCE_KEYS == fe_listed
    assert mod._SOE_SOURCE_KEYS == fe_soe


def test_dr_row_defs_mirror_frontend():
    src = _read("f2DataResourceInventory.ts")
    start = src.index("const ROW_DEFS")
    end = src.index("\n]", start)
    block = src[start:end]
    fe = [
        (m.group(1), m.group(2), m.group(3))
        for m in re.finditer(
            r"rowKey:\s*'([^']+)',\s*label:\s*'([^']+)',\s*kind:\s*'([^']+)'", block
        )
    ]
    assert len(fe) == 21, f"前端 DR 骨架应为 21 行，实为 {len(fe)}"
    assert list(mod._DR_ROW_DEFS) == fe


def test_dr_col_labels_mirror_frontend():
    src = _read("f2DataResourceInventory.ts")
    for key, label in mod._DR_COL_LABELS.items():
        assert f"{key}: '{label}'" in src, f"DR 列名 {key} 与前端不一致"


def test_dr_soe_label_override_mirror_frontend():
    src = _read("f2DataResourceInventory.ts")
    for key, label in mod._DR_SOE_LABEL_OVERRIDES.items():
        assert f"'{key}': '{label}'" in src


# ════════════════════════════════════════════════════════════════════
# 2. workbook 结构
# ════════════════════════════════════════════════════════════════════


@pytest.mark.parametrize("variant,expect_blocks", [("listed", 7), ("soe", 2)])
def test_template_sheets(variant: str, expect_blocks: int):
    wb = mod.build_disclosure_workbook(variant)  # type: ignore[arg-type]
    names = wb.sheetnames
    assert names[0] == mod._GUIDE_SHEET
    assert mod._CLASS_SHEET in names
    assert mod._TEXT_SHEET in names
    for blk in mod.blocks_of(variant):  # type: ignore[arg-type]
        assert blk.sheet in names, f"{variant} 缺区块 sheet {blk.sheet}"
    assert len(mod.blocks_of(variant)) == expect_blocks  # type: ignore[arg-type]
    # sheet 名必须满足 Excel 限制
    for n in names:
        assert len(n) <= 31, f"sheet 名超 31 字符: {n}"
        assert not set(n) & set("[]:*?/\\"), f"sheet 名含非法字符: {n}"


@pytest.mark.parametrize("variant", ["listed", "soe"])
def test_template_headers_and_labels(variant: str):
    wb = mod.build_disclosure_workbook(variant)  # type: ignore[arg-type]
    for blk in mod.blocks_of(variant):  # type: ignore[arg-type]
        ws = wb[blk.sheet]
        hdr_row = mod._find_header_row(ws, blk.headers[0])
        assert hdr_row is not None, f"{blk.sheet} 找不到表头行"
        actual = [
            str(ws.cell(row=hdr_row, column=c).value or "")
            for c in range(1, len(blk.headers) + 1)
        ]
        assert actual == list(blk.headers)
        if blk.kind == "override":
            labels = [
                ws.cell(row=r, column=1).value
                for r in range(hdr_row + 1, hdr_row + 1 + len(mod.categories_of(variant)))  # type: ignore[arg-type]
            ]
            assert labels == [lb for _, lb in mod.categories_of(variant)]  # type: ignore[arg-type]
        if blk.kind == "dr":
            keys = [
                ws.cell(row=r, column=1).value
                for r in range(hdr_row + 1, hdr_row + 1 + 21)
            ]
            labels = [
                ws.cell(row=r, column=2).value
                for r in range(hdr_row + 1, hdr_row + 1 + 21)
            ]
            assert keys == [rk for rk, _, _ in mod.dr_rows_of(variant)]  # type: ignore[arg-type]
            assert labels == [lb for _, lb, _ in mod.dr_rows_of(variant)]  # type: ignore[arg-type]


def test_soe_dr_section_label_differs():
    wb_l = mod.build_disclosure_workbook("listed")
    wb_s = mod.build_disclosure_workbook("soe")
    col_l = [c.value for c in wb_l["(8)数据资源"]["B"]]
    col_s = [c.value for c in wb_s["(5)数据资源"]["B"]]
    assert "二、存货跌价准备" in col_l
    assert "二、跌价准备" in col_s
    assert "二、存货跌价准备" not in col_s


@pytest.mark.parametrize("variant", ["listed", "soe"])
def test_text_sheet_lists_all_note_blocks(variant: str):
    wb = mod.build_disclosure_workbook(variant)  # type: ignore[arg-type]
    ws = wb[mod._TEXT_SHEET]
    names = [ws.cell(row=r, column=1).value for r in range(3, ws.max_row + 1)]
    assert names == [n for _, n in mod.texts_of(variant)]  # type: ignore[arg-type]


@pytest.mark.parametrize("variant", ["listed", "soe"])
def test_item_id_prefix_matches_frontend(variant: str):
    """item_id 必须与前端 PREFIX 一致，否则导入写到孤儿键。"""
    fname = "useF2DisclosureListed.ts" if variant == "listed" else "useF2DisclosureSoe.ts"
    src = _read(fname)
    assert f"const PREFIX = 'F2-note-{variant}-'" in src
    for blk in mod.blocks_of(variant):  # type: ignore[arg-type]
        assert f"`${{PREFIX}}{blk.suffix}`" in src, f"{blk.suffix} 不在前端 item 列表中"
    for suffix, _ in mod.texts_of(variant):  # type: ignore[arg-type]
        assert f"`${{PREFIX}}{suffix}`" in src, f"文本 {suffix} 不在前端 item 列表中"


# ════════════════════════════════════════════════════════════════════
# 3. 导出 → 导入 往返（纯解析层，不碰 DB）
# ════════════════════════════════════════════════════════════════════


def _wb_bytes(wb) -> bytes:
    buf = io.BytesIO()
    wb.save(buf)
    return buf.getvalue()


def _reload(wb):
    return load_workbook(io.BytesIO(_wb_bytes(wb)), data_only=True)


def _set(ws, header_first: str, row_label: str, col_name: str, value):
    hdr = mod._find_header_row(ws, header_first)
    idx = mod._header_index(ws, hdr, 12)
    for r in range(hdr + 1, ws.max_row + 1):
        if str(ws.cell(row=r, column=1).value or "").strip() == row_label:
            ws.cell(row=r, column=idx[col_name]).value = value
            return r
    raise AssertionError(f"未找到行 {row_label}")


def test_roundtrip_override_listed():
    wb = mod.build_disclosure_workbook("listed")
    blk = next(b for b in mod._LISTED_BLOCKS if b.kind == "override")
    ws = wb[blk.sheet]
    _set(ws, blk.headers[0], "原材料", "本期计提", 1234.5)
    _set(ws, blk.headers[0], "库存商品", "本期转回", 200)

    wb2 = _reload(wb)
    store, warns, cnt = mod._import_override(wb2[blk.sheet], blk, "listed")
    assert cnt == 2
    assert store["raw-materials"] == {"incProvision": 1234.5}
    assert store["finished-goods"] == {"decReversal": 200.0}
    # 未填的类别不出现（空 = 不覆盖，沿用 F2-1 自动取数）
    assert "data-resources" not in store
    assert warns == []


def test_override_unknown_category_warns_and_skips():
    blk = next(b for b in mod._LISTED_BLOCKS if b.kind == "override")
    wb = mod.build_disclosure_workbook("listed")
    ws = wb[blk.sheet]
    hdr = mod._find_header_row(ws, blk.headers[0])
    ws.cell(row=ws.max_row + 1, column=1).value = "不存在的类别"
    ws.cell(row=ws.max_row, column=2).value = 99

    store, warns, cnt = mod._import_override(_reload(wb)[blk.sheet], blk, "listed")
    assert cnt == 0
    assert store == {}
    assert any("不存在的类别" in w for w in warns)


def test_soe_override_has_writeoff_column():
    blk = next(b for b in mod._SOE_BLOCKS if b.kind == "override")
    assert "decWriteOff" in blk.keys, "国企跌价准备变动应含「本期核销」"
    assert "本期核销" in blk.headers
    wb = mod.build_disclosure_workbook("soe")
    ws = wb[blk.sheet]
    _set(ws, blk.headers[0], "原材料", "本期核销", 55)
    store, _w, cnt = mod._import_override(_reload(wb)[blk.sheet], blk, "soe")
    assert cnt == 1
    assert store["raw-combined"] == {"decWriteOff": 55.0}


def test_listed_override_has_no_writeoff_column():
    """上市源模板无「本期核销」列，不得凭空加。"""
    blk = next(b for b in mod._LISTED_BLOCKS if b.kind == "override")
    assert "decWriteOff" not in blk.keys
    assert "本期核销" not in blk.headers


def test_roundtrip_rows_listed_dev_cost():
    blk = next(b for b in mod._LISTED_BLOCKS if b.suffix == "s5-rows")
    wb = mod.build_disclosure_workbook("listed")
    ws = wb[blk.sheet]
    hdr = mod._find_header_row(ws, blk.headers[0])
    idx = mod._header_index(ws, hdr, len(blk.headers))
    r = hdr + 1
    ws.cell(row=r, column=idx["项目名称"]).value = "A 项目"
    ws.cell(row=r, column=idx["开工时间"]).value = "2024-03"
    ws.cell(row=r, column=idx["预计总投资"]).value = 5000000
    ws.cell(row=r, column=idx["期末余额"]).value = 1200000.55

    rows, warns, cnt = mod._import_rows(_reload(wb)[blk.sheet], blk)
    assert cnt == 1
    row = rows[0]
    assert row["projectName"] == "A 项目"
    assert row["startDate"] == "2024-03"
    assert row["estimatedInvestment"] == 5000000.0
    assert row["endBalance"] == 1200000.55
    assert row["priorBalance"] == 0.0  # 空数值列 → 0
    assert row["rowId"].startswith("s5-rows-")
    assert warns == []


def test_rows_skips_fully_blank_rows():
    blk = next(b for b in mod._LISTED_BLOCKS if b.suffix == "s7-rows")
    wb = mod.build_disclosure_workbook("listed")
    rows, _w, cnt = mod._import_rows(_reload(wb)[blk.sheet], blk)
    assert cnt == 0 and rows == []  # 模板的 3 行占位全空 → 不产生垃圾行


def test_s3_prior_and_end_are_separate_blocks():
    ends = [b for b in mod._LISTED_BLOCKS if b.suffix == "s3-end"]
    priors = [b for b in mod._LISTED_BLOCKS if b.suffix == "s3-prior"]
    assert len(ends) == 1 and len(priors) == 1
    assert ends[0].sheet != priors[0].sheet
    assert ends[0].headers == priors[0].headers


def test_s3_excludes_formula_columns():
    """净值/占比/计提比例是公式列，不得进导入列。"""
    blk = next(b for b in mod._LISTED_BLOCKS if b.suffix == "s3-end")
    for forbidden in ("netValue", "balancePct", "impairmentPct"):
        assert forbidden not in blk.keys
    assert "账面价值" not in blk.headers
    assert "计提比例" not in blk.headers


@pytest.mark.parametrize("variant", ["listed", "soe"])
def test_roundtrip_dr(variant: str):
    """DR 按 rowKey 匹配：账面原值段与跌价准备段的「1.期初余额」必须互不干扰。"""
    wb = mod.build_disclosure_workbook(variant)  # type: ignore[arg-type]
    blk = next(b for b in mod.blocks_of(variant) if b.kind == "dr")  # type: ignore[arg-type]
    ws = wb[blk.sheet]
    _set(ws, blk.headers[0], "gross-open", mod._DR_COL_LABELS["purchased"], 100)
    _set(ws, blk.headers[0], "imp-open", mod._DR_COL_LABELS["purchased"], 30)
    _set(ws, blk.headers[0], "gross-inc", mod._DR_COL_LABELS["selfProcessed"], 250)

    store, warns, cnt = mod._import_dr(_reload(wb)[blk.sheet], blk, variant)  # type: ignore[arg-type]
    assert cnt == 3
    assert store["gross-open"] == {"purchased": 100.0}
    assert store["imp-open"] == {"purchased": 30.0}  # 同名标签不再互相覆盖
    assert store["gross-inc"] == {"selfProcessed": 250.0}
    assert warns == []


@pytest.mark.parametrize("variant", ["listed", "soe"])
def test_dr_duplicate_labels_exist_hence_key_column_required(variant: str):
    """守卫：DR 骨架确实存在重复标签 → 「行标识」列不可省。"""
    labels = [lb for _rk, lb, _k in mod.dr_rows_of(variant)]  # type: ignore[arg-type]
    dupes = {lb for lb in labels if labels.count(lb) > 1}
    assert dupes, "若骨架标签变唯一，可考虑简化为按标签匹配"
    blk = next(b for b in mod.blocks_of(variant) if b.kind == "dr")  # type: ignore[arg-type]
    assert blk.headers[0] == mod._DR_KEY_HEADER
    assert blk.headers[1] == mod._DR_LABEL_HEADER


@pytest.mark.parametrize("variant", ["listed", "soe"])
def test_dr_ignores_section_and_derived_rows(variant: str):
    """段标题行与公式行即使被填了也不导入。"""
    wb = mod.build_disclosure_workbook(variant)  # type: ignore[arg-type]
    blk = next(b for b in mod.blocks_of(variant) if b.kind == "dr")  # type: ignore[arg-type]
    ws = wb[blk.sheet]
    _set(ws, blk.headers[0], "gross-end", mod._DR_COL_LABELS["purchased"], 999)
    _set(ws, blk.headers[0], "gross-section", mod._DR_COL_LABELS["purchased"], 888)

    store, _w, _c = mod._import_dr(_reload(wb)[blk.sheet], blk, variant)  # type: ignore[arg-type]
    assert "gross-end" not in store
    assert "gross-section" not in store
    assert all(v.get("purchased") not in (999.0, 888.0) for v in store.values())


@pytest.mark.parametrize("variant", ["listed", "soe"])
def test_dr_unknown_row_key_warns(variant: str):
    wb = mod.build_disclosure_workbook(variant)  # type: ignore[arg-type]
    blk = next(b for b in mod.blocks_of(variant) if b.kind == "dr")  # type: ignore[arg-type]
    ws = wb[blk.sheet]
    ws.cell(row=ws.max_row + 1, column=1).value = "made-up-key"
    ws.cell(row=ws.max_row, column=3).value = 1
    _set(ws, blk.headers[0], "gross-open", mod._DR_COL_LABELS["purchased"], 5)

    store, warns, cnt = mod._import_dr(_reload(wb)[blk.sheet], blk, variant)  # type: ignore[arg-type]
    assert cnt == 1
    assert "made-up-key" not in store
    assert any("made-up-key" in w for w in warns)


@pytest.mark.parametrize("variant", ["listed", "soe"])
def test_roundtrip_texts(variant: str):
    wb = mod.build_disclosure_workbook(variant)  # type: ignore[arg-type]
    ws = wb[mod._TEXT_SHEET]
    first_suffix, first_name = mod.texts_of(variant)[0]  # type: ignore[arg-type]
    _set(ws, mod._TEXT_HEADERS[0], first_name, mod._TEXT_HEADERS[1], "多行\n文本内容")

    texts, warns, cnt = mod._import_texts(_reload(wb)[mod._TEXT_SHEET], variant)  # type: ignore[arg-type]
    assert cnt == len(mod.texts_of(variant))  # type: ignore[arg-type]
    assert texts[first_suffix] == "多行\n文本内容"
    assert warns == []


def test_export_with_data_renders_values():
    """带数据导出：override / rows / dr / text 都写出可读值。"""
    remarks = {
        "F2-note-listed-s2-overrides": json.dumps({"raw-materials": {"incProvision": 88.5}}),
        "F2-note-listed-s5-rows": json.dumps(
            [{"rowId": "x", "projectName": "B 项目", "endBalance": 777.25}]
        ),
        "F2-note-listed-s8-data-resource": json.dumps({"gross-open": {"purchased": 42}}),
        "F2-note-listed-note-borrow": "本期资本化借款费用 12 万元。",
    }
    wb = mod.build_disclosure_workbook("listed", remarks=remarks, class_rows=[])
    wb2 = _reload(wb)

    ov = mod._import_override(
        wb2["(2)跌价准备变动"],
        next(b for b in mod._LISTED_BLOCKS if b.kind == "override"),
        "listed",
    )[0]
    assert ov["raw-materials"]["incProvision"] == 88.5

    rows = mod._import_rows(
        wb2["(5)开发成本"], next(b for b in mod._LISTED_BLOCKS if b.suffix == "s5-rows")
    )[0]
    assert rows[0]["projectName"] == "B 项目"
    assert rows[0]["endBalance"] == 777.25

    dr = mod._import_dr(
        wb2["(8)数据资源"], next(b for b in mod._LISTED_BLOCKS if b.kind == "dr"), "listed"
    )[0]
    assert dr["gross-open"]["purchased"] == 42.0

    txt = mod._import_texts(wb2[mod._TEXT_SHEET], "listed")[0]
    assert txt["note-borrow"] == "本期资本化借款费用 12 万元。"


def test_malformed_remark_json_does_not_crash_export():
    remarks = {
        "F2-note-soe-s2-overrides": "{ not json",
        "F2-note-soe-s5-data-resource": "[]",  # 类型不符（期望 dict）
    }
    wb = mod.build_disclosure_workbook("soe", remarks=remarks, class_rows=[])
    assert mod._TEXT_SHEET in wb.sheetnames


def test_missing_sheet_is_skipped_not_error():
    """缺 sheet 应跳过而非报错（用户只想导入部分区块）。"""
    from openpyxl import Workbook as _Wb

    wb = _Wb()
    wb.active.title = mod._TEXT_SHEET
    ws = wb.active
    ws.append(["提示"])
    ws.append(list(mod._TEXT_HEADERS))
    ws.append([mod.texts_of("listed")[0][1], "仅导入文本"])
    texts, _w, cnt = mod._import_texts(_reload(wb)[mod._TEXT_SHEET], "listed")
    assert cnt == 1
    assert texts[mod.texts_of("listed")[0][0]] == "仅导入文本"


def test_header_row_tolerates_reordered_columns():
    """列被用户拖动顺序后仍按名匹配。"""
    blk = next(b for b in mod._LISTED_BLOCKS if b.suffix == "s7-rows")
    from openpyxl import Workbook as _Wb

    wb = _Wb()
    ws = wb.active
    ws.title = blk.sheet
    ws.append(["提示：xxx"])
    ws.append(["项目名称", "期末余额", "期初余额", "本期增加", "本期减少"])
    ws.append(["C 项目", 500, 100, 400, 0])
    rows, _w, cnt = mod._import_rows(_reload(wb)[blk.sheet], blk)
    assert cnt == 1
    assert rows[0]["opening"] == 100.0
    assert rows[0]["ending"] == 500.0


def test_override_is_whole_table_replace_clearing_a_row_revokes_it():
    """整表覆盖：清空某行 = 撤销该类别覆盖（该 rowKey 不出现在新 store 中）。"""
    blk = next(b for b in mod._LISTED_BLOCKS if b.kind == "override")
    wb = mod.build_disclosure_workbook(
        "listed",
        remarks={
            "F2-note-listed-s2-overrides": json.dumps(
                {"raw-materials": {"incProvision": 10}, "finished-goods": {"incProvision": 20}}
            )
        },
        class_rows=[],
    )
    ws = wb[blk.sheet]
    # 模拟用户把「原材料」那格清空，只保留「库存商品」
    _set(ws, blk.headers[0], "原材料", "本期计提", None)

    store, _w, filled = mod._import_override(_reload(wb)[blk.sheet], blk, "listed")
    assert filled == 1
    assert "raw-materials" not in store, "清空的行应撤销覆盖"
    assert store["finished-goods"] == {"incProvision": 20.0}


def test_override_all_blank_returns_zero_so_caller_skips():
    """全空模板不得清掉既有覆盖 → filled=0，上层据此跳过写库。"""
    blk = next(b for b in mod._LISTED_BLOCKS if b.kind == "override")
    wb = mod.build_disclosure_workbook("listed")  # 模板：全空
    store, _w, filled = mod._import_override(_reload(wb)[blk.sheet], blk, "listed")
    assert filled == 0
    assert store == {}


def test_override_note_documents_whole_table_semantics():
    """文档与实现一致性守卫（曾出现 note 写「留空=沿用自动取数」但语义未说明整表覆盖）。"""
    for blocks in (mod._LISTED_BLOCKS, mod._SOE_BLOCKS):
        blk = next(b for b in blocks if b.kind == "override")
        assert "整表覆盖" in blk.note
        assert "导出数据" in blk.note


def test_dr_is_also_whole_table_replace():
    blk = next(b for b in mod._LISTED_BLOCKS if b.kind == "dr")
    wb = mod.build_disclosure_workbook(
        "listed",
        remarks={
            "F2-note-listed-s8-data-resource": json.dumps(
                {"gross-open": {"purchased": 1}, "imp-open": {"purchased": 2}}
            )
        },
        class_rows=[],
    )
    ws = wb[blk.sheet]
    _set(ws, blk.headers[0], "gross-open", mod._DR_COL_LABELS["purchased"], None)
    store, _w, cnt = mod._import_dr(_reload(wb)[blk.sheet], blk, "listed")
    assert cnt == 1
    assert "gross-open" not in store
    assert store["imp-open"] == {"purchased": 2.0}


def test_class_sheet_is_export_only():
    """(1) 分类表不在可导入区块中（跨表取数，导入会与 F2-1 打架）。"""
    for variant in ("listed", "soe"):
        assert all(
            b.sheet != mod._CLASS_SHEET for b in mod.blocks_of(variant)  # type: ignore[arg-type]
        )


@pytest.mark.parametrize("variant", ["listed", "soe"])
def test_guidance_mentions_readonly_class_sheet(variant: str):
    lines = "\n".join(mod._guidance_lines(variant))  # type: ignore[arg-type]
    assert mod._CLASS_SHEET in lines
    assert "导入时忽略" in lines

"""F1 公式预设守卫（`prefill_formula_mapping.json` → 公式管理页 `workpaper:F1`）。

背景：改造前 `workpaper:F1` 的 17 条预设全部属 `审定表F1-1` / `明细表F1-2` /
`实质性分析F1-4` 三个 sheet —— **两个披露 sheet 与 F1-5 在公式管理页一片空白**，
而源模板这三页的每个数据格都是跨 sheet 公式
（`审定表F1-1'!I17` / `实质性分析F1-4'!B41` / `长期挂款检查表F1-5'!J6`）。

本守卫钉住三件事：

1. **`cell_ref` 页内唯一** —— `convert_prefill_presets` 的 `page_key` 是
   `workpaper:{wp_code}`、**忽略 sheet**，同名 `cell_ref` 会被静默去重丢弃
   （平台现存 `上年审定数` 在 25+ 循环内撞键就是这个坑）。
2. **三个新 sheet 齐备**，且 sheet 名与源 xlsx tab 名逐字一致。
3. **新增公式语法合法**（`validate_formula` 返回空错误列表）且**披露块不被 `WP()` 引用**
   —— 披露块只允许作为 `WP()` 的消费方，反向引用会形成循环。

spec: .kiro/specs/f1-extraction-chain-and-disclosure-source-fidelity/ (Task 4.2)
"""
from __future__ import annotations

import json
import re
from pathlib import Path

import openpyxl
import pytest

from app.services.formula_engine import validate_formula
from app.services.formula_management.preset_library import convert_prefill_presets

_ROOT = Path(__file__).resolve().parent.parent
_MAPPING = _ROOT / "data" / "prefill_formula_mapping.json"
_SRC_XLSX = _ROOT / "wp_templates" / "F" / "F1 预付账款.xlsx"

PAGE_KEY = "workpaper:F1"

SHEET_ADJ = "审定表F1-1"
SHEET_DETAIL = "明细表F1-2"
SHEET_ANALYSIS = "实质性分析F1-4"
SHEET_LT = "长期挂款检查表F1-5"
SHEET_DISC_LISTED = "附注披露信息(上市公司)"
SHEET_DISC_SOE = "附注披露信息(国企)"

#: 本 spec 新增的三个 sheet（改造前公式管理页在这三页一片空白）
_NEW_SHEETS = (SHEET_LT, SHEET_DISC_LISTED, SHEET_DISC_SOE)

#: 已存在于平台但**未注册进 `formula_engine._REGISTRY`** 的 prefill 专属词汇。
#: 它们是改造前的既有条目，不属本 spec 范围 → 语法校验时豁免。
#: 反向自检：若某天被注册了，本豁免必须移除（见 `test_prefill_only_vocab_still_unregistered`）。
_PREFILL_ONLY_FUNCS = ("ADJ", "TB_SUM")


def _f1_blocks() -> list[dict]:
    data = json.loads(_MAPPING.read_text(encoding="utf-8"))
    return [b for b in data["mappings"] if b.get("wp_code") == "F1"]


def _f1_presets() -> list:
    return [p for p in convert_prefill_presets() if getattr(p, "page_key", "") == PAGE_KEY]


def _uses_prefill_only_func(formula: str) -> bool:
    return any(f"{fn}(" in formula for fn in _PREFILL_ONLY_FUNCS)


# ─────────────────── Property 8：cell_ref 页内唯一 ───────────────────


def test_f1_cell_refs_unique_within_page():
    """`page_key` 忽略 sheet → `cell_ref` 必须页内唯一，否则被静默去重丢弃。"""
    refs = [c["cell_ref"] for b in _f1_blocks() for c in b["cells"]]
    dupes = sorted({r for r in refs if refs.count(r) > 1})
    assert dupes == [], f"workpaper:F1 内 cell_ref 撞键（会被静默丢弃）：{dupes}"


def test_runtime_presets_not_silently_deduped():
    """运行态条目数 == 唯一 target_cell 数（真正验证「没被吞」）。"""
    presets = _f1_presets()
    cells = [p.target_cell for p in presets]
    assert len(cells) == len(set(cells)), sorted(
        {c for c in cells if cells.count(c) > 1}
    )
    # JSON 里声明的条数必须全部到达运行态
    declared = sum(len(b["cells"]) for b in _f1_blocks())
    assert len(presets) == declared, f"declared={declared} runtime={len(presets)}"


def test_runtime_preset_count_grew():
    """改造前 17 条（三个 sheet）→ 现在必须覆盖 6 个 sheet 且条数更多。"""
    assert len(_f1_presets()) > 17
    assert {b["sheet"] for b in _f1_blocks()} == {
        SHEET_ADJ, SHEET_DETAIL, SHEET_ANALYSIS,
        SHEET_LT, SHEET_DISC_LISTED, SHEET_DISC_SOE,
    }


# ─────────────────── sheet 名与源 xlsx tab 名逐字一致 ───────────────────


def test_preset_sheet_names_match_source_xlsx_tabs():
    """所有 F1 预设块的 sheet 名必须是源 xlsx 的真实 tab 名（逐字，含括号宽度）。"""
    tabs = set(openpyxl.load_workbook(_SRC_XLSX, read_only=True).sheetnames)
    for block in _f1_blocks():
        assert block["sheet"] in tabs, (
            f"预设 sheet {block['sheet']!r} 不在源 xlsx tab 名里：{sorted(tabs)}"
        )


def test_disclosure_sheet_names_use_halfwidth_parens():
    """F1 源模板披露 tab 用**半角**括号（与 G/H/K 系全角不同，别「顺手改成全角」）。"""
    assert SHEET_DISC_LISTED == "附注披露信息(上市公司)"
    assert SHEET_DISC_SOE == "附注披露信息(国企)"
    tabs = set(openpyxl.load_workbook(_SRC_XLSX, read_only=True).sheetnames)
    assert SHEET_DISC_LISTED in tabs and SHEET_DISC_SOE in tabs


# ─────────────────── 新增 sheet 的内容完备性 ───────────────────


@pytest.mark.parametrize("sheet", _NEW_SHEETS)
def test_new_sheet_has_presets(sheet: str):
    block = next((b for b in _f1_blocks() if b["sheet"] == sheet), None)
    assert block is not None, f"缺少 sheet={sheet} 的预设块"
    assert len(block["cells"]) >= 3, block["cells"]
    for cell in block["cells"]:
        assert cell["cell_ref"].strip()
        assert cell["formula"].startswith("=")
        assert len(str(cell.get("description") or "")) >= 10, cell["cell_ref"]


def test_soe_disclosure_has_long_term_sheet_reference():
    """国企②表的取数来源必须显式指向 F1-5（源模板 C18 = F1-5!J6 的口径）。"""
    block = next(b for b in _f1_blocks() if b["sheet"] == SHEET_DISC_SOE)
    formulas = [c["formula"] for c in block["cells"]]
    assert any(SHEET_LT in f for f in formulas), formulas
    # 且描述里要写明是**审定余额**（已扣坏账），防后人误改成期末余额毛额
    hit = next(c for c in block["cells"] if SHEET_LT in c["formula"])
    assert "审定余额" in hit["description"]


def test_listed_disclosure_top5_references_analysis_sheet():
    """上市③前五名来源 = F1-4（源模板 A28:B32 = 实质性分析F1-4!A41:B45）。"""
    block = next(b for b in _f1_blocks() if b["sheet"] == SHEET_DISC_LISTED)
    assert any(SHEET_ANALYSIS in c["formula"] for c in block["cells"])


def test_provision_uses_subdivided_standard_code():
    """备抵一律用细分标准码 `1231-04`；写宽口径 `1231` 会含应收账款坏账。"""
    for block in _f1_blocks():
        for cell in block["cells"]:
            assert not re.search(r"TB\('1231'", cell["formula"]), (
                f"{block['sheet']}/{cell['cell_ref']} 用了宽口径 1231：{cell['formula']}"
            )


# ─────────────────── 语法合法 + 防循环 ───────────────────


def test_all_f1_formulas_are_syntactically_valid():
    """除 prefill 专属词汇（ADJ/TB_SUM）外，全部公式语法合法。"""
    bad: list[str] = []
    for block in _f1_blocks():
        for cell in block["cells"]:
            formula = cell["formula"]
            if _uses_prefill_only_func(formula):
                continue
            errs = validate_formula(formula)
            if errs:
                bad.append(f"{block['sheet']}/{cell['cell_ref']}: {formula} → {errs}")
    assert bad == [], bad


def test_prefill_only_vocab_still_unregistered():
    """反向自检：ADJ/TB_SUM 若被注册进公式引擎，上面的豁免就该删掉。"""
    for fn in _PREFILL_ONLY_FUNCS:
        errs = validate_formula(f"=({fn}('1123','x'))")
        assert errs, (
            f"{fn}() 现已被 formula_engine 识别 → 请从 _PREFILL_ONLY_FUNCS 移除豁免"
        )


def test_disclosure_blocks_are_not_referenced_by_wp():
    """披露块只能作 `WP()` 的消费方，不得被任何块 `WP()` 引用（防循环）。"""
    data = json.loads(_MAPPING.read_text(encoding="utf-8"))
    for block in data["mappings"]:
        for cell in block.get("cells") or []:
            formula = str(cell.get("formula") or "")
            for sheet in (SHEET_DISC_LISTED, SHEET_DISC_SOE):
                assert f"'{sheet}'" not in formula, (
                    f"{block.get('wp_code')}/{block.get('sheet')}/{cell.get('cell_ref')} "
                    f"引用了披露 sheet {sheet}：{formula}"
                )


def test_detail_sheet_has_no_wp_reference():
    """明细表 F1-2 是取数链路的根，禁 `WP()`（否则与审定表/披露表成环）。"""
    block = next(b for b in _f1_blocks() if b["sheet"] == SHEET_DETAIL)
    for cell in block["cells"]:
        assert "WP(" not in cell["formula"], cell

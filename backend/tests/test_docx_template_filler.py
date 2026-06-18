"""Word 导出引擎单元测试（PRE-2，docx_template_filler）.

用 python-docx 构造含红/蓝/黑 run + 注释表的内存文档，验证颜色语义：
- 红色占位符替换 + 转黑
- 蓝色【】编制提示 run 删除
- 黑色固定正文保留
- 注释表（关键词识别）整张删除
- 占位符在 table 单元格内也替换

颜色语义只在 filler 实现；编排层禁止重复实现（见 requirements PRE-2）。
"""

from __future__ import annotations

import pytest

docx = pytest.importorskip("docx")

from docx import Document
from docx.shared import RGBColor

from app.services.docx_template_filler import (
    ColorSemanticsConfig,
    FillResult,
    build_replacements,
    export_to_bytes,
    fill_and_export,
    is_blue_run,
    is_note_table,
    is_red_run,
)

# 颜色常量
RED = RGBColor(0xFF, 0x00, 0x00)       # red>200, green<100 → 红
BLUE = RGBColor(0x00, 0x00, 0xFF)      # red<100, blue>150 → 蓝
BLACK = RGBColor(0x00, 0x00, 0x00)     # 黑


def _add_colored_run(paragraph, text, color):
    run = paragraph.add_run(text)
    if color is not None:
        run.font.color.rgb = color
    return run


# ─── Fixtures ─────────────────────────────────────────────────────────────────


@pytest.fixture
def context():
    return {"client_name": "测试科技有限公司", "audit_year": "2025"}


@pytest.fixture
def sample_doc():
    """构造含红/蓝/黑 run + 普通表 + 注释表的文档."""
    doc = Document()

    # 段落 1：红色占位符（公司名 + 年度）
    p1 = doc.add_paragraph()
    _add_colored_run(p1, "××公司", RED)
    _add_colored_run(p1, "202X年度审计", RED)

    # 段落 2：蓝色编制提示（应删除）+ 黑色正文（应保留）
    p2 = doc.add_paragraph()
    _add_colored_run(p2, "【请填写公司名称】", BLUE)
    _add_colored_run(p2, "固定正文保留", BLACK)

    # 普通表（含占位符，应替换，不删除）
    t1 = doc.add_table(rows=1, cols=2)
    t1.rows[0].cells[0].text = "项目"
    cell_p = t1.rows[0].cells[1].paragraphs[0]
    _add_colored_run(cell_p, "××公司", RED)

    # 注释表（表首单元格含「注：」，应整张删除）
    t2 = doc.add_table(rows=1, cols=1)
    t2.rows[0].cells[0].text = "注：本表为模板参考说明，导出时删除"

    return doc


# ─── 颜色判定 ────────────────────────────────────────────────────────────────


def test_is_red_run():
    doc = Document()
    p = doc.add_paragraph()
    assert is_red_run(_add_colored_run(p, "x", RED)) is True
    assert is_red_run(_add_colored_run(p, "x", BLUE)) is False
    assert is_red_run(_add_colored_run(p, "x", BLACK)) is False
    assert is_red_run(_add_colored_run(p, "x", None)) is False


def test_is_blue_run():
    doc = Document()
    p = doc.add_paragraph()
    assert is_blue_run(_add_colored_run(p, "x", BLUE)) is True
    assert is_blue_run(_add_colored_run(p, "x", RED)) is False
    assert is_blue_run(_add_colored_run(p, "x", BLACK)) is False
    assert is_blue_run(_add_colored_run(p, "x", None)) is False


# ─── 占位符映射 ──────────────────────────────────────────────────────────────


def test_build_replacements_basic():
    r = build_replacements({"client_name": "甲公司", "audit_year": "2025"})
    assert r["××公司"] == "甲公司"
    assert r["XX公司"] == "甲公司"
    assert r["ABC公司"] == "甲公司"
    assert r["202X"] == "2025"
    assert r["201X"] == "2024"  # 上年度自动推算


def test_build_replacements_defaults():
    r = build_replacements({})
    assert r["××公司"] == "XX公司"
    assert r["202X"] == "202X"
    assert r["201X"] == "201X"


def test_build_replacements_extra_placeholders():
    r = build_replacements({"placeholders": {"{{partner}}": "张三", "{{empty}}": None}})
    assert r["{{partner}}"] == "张三"
    assert r["{{empty}}"] == ""


# ─── 核心 fill_and_export ────────────────────────────────────────────────────


def test_red_placeholder_replaced_and_black(sample_doc, context):
    doc, result = fill_and_export(sample_doc, context)

    p1 = doc.paragraphs[0]
    # 红色占位符被替换为项目数据
    assert "测试科技有限公司" in p1.text
    assert "2025" in p1.text
    assert "××公司" not in p1.text
    assert "202X" not in p1.text
    # 替换后转黑
    for run in p1.runs:
        rgb = run.font.color.rgb
        assert rgb == RGBColor(0, 0, 0)


def test_blue_guidance_run_deleted(sample_doc, context):
    doc, _ = fill_and_export(sample_doc, context)
    p2 = doc.paragraphs[1]
    # 蓝色【】编制提示删除，黑色正文保留
    assert "【请填写公司名称】" not in p2.text
    assert "固定正文保留" in p2.text


def test_black_body_preserved(sample_doc, context):
    doc, _ = fill_and_export(sample_doc, context)
    full_text = "\n".join(p.text for p in doc.paragraphs)
    assert "固定正文保留" in full_text


def test_note_table_deleted(sample_doc, context):
    assert len(sample_doc.tables) == 2
    doc, result = fill_and_export(sample_doc, context)
    # 注释表删除，普通表保留
    assert result.note_tables_deleted == 1
    assert len(doc.tables) == 1


def test_placeholder_replaced_in_table_cell(sample_doc, context):
    doc, _ = fill_and_export(sample_doc, context)
    # 剩余唯一的普通表，单元格内占位符也被替换
    remaining = doc.tables[0]
    cell_text = remaining.rows[0].cells[1].text
    assert "测试科技有限公司" in cell_text
    assert "××公司" not in cell_text


def test_export_to_bytes_roundtrip(sample_doc, context):
    doc, _ = fill_and_export(sample_doc, context)
    data = export_to_bytes(doc)
    assert isinstance(data, bytes)
    assert len(data) > 0
    # 字节流可被重新打开
    import io
    reopened = Document(io.BytesIO(data))
    assert reopened is not None


# ─── 可配置注释表策略 ────────────────────────────────────────────────────────


def test_note_table_custom_keywords():
    doc = Document()
    t = doc.add_table(rows=1, cols=1)
    t.rows[0].cells[0].text = "填表说明：仅供参考"
    cfg = ColorSemanticsConfig(note_table_keywords=("填表说明",))
    assert is_note_table(t, cfg) is True
    assert is_note_table(t, ColorSemanticsConfig(note_table_keywords=("不存在",))) is False


def test_disable_note_table_deletion(sample_doc, context):
    cfg = ColorSemanticsConfig(delete_note_tables=False)
    doc, result = fill_and_export(sample_doc, context, cfg)
    assert result.note_tables_deleted == 0
    assert len(doc.tables) == 2  # 注释表保留


def test_disable_blue_deletion(context):
    doc = Document()
    p = doc.add_paragraph()
    _add_colored_run(p, "【提示】", BLUE)
    cfg = ColorSemanticsConfig(delete_blue_runs=False)
    out, _ = fill_and_export(doc, context, cfg)
    assert "【提示】" in out.paragraphs[0].text


# ─── Property-based tests (hypothesis, max_examples=5) ────────────────────────

from hypothesis import given, settings
from hypothesis import strategies as st


@settings(max_examples=5, deadline=None)
@given(
    client=st.text(min_size=1, max_size=20).filter(lambda s: "×" not in s and "X" not in s),
    year=st.integers(min_value=2000, max_value=2099),
)
def test_property_red_placeholder_always_replaced(client, year):
    """任意公司名/年度：红色占位符总被替换且无残留 + 转黑."""
    doc = Document()
    p = doc.add_paragraph()
    _add_colored_run(p, "××公司202X年度", RED)
    out, _ = fill_and_export(doc, {"client_name": client, "audit_year": str(year)})
    text = out.paragraphs[0].text
    assert "××公司" not in text
    assert "202X" not in text
    assert client in text
    assert str(year) in text
    for run in out.paragraphs[0].runs:
        assert run.font.color.rgb == RGBColor(0, 0, 0)


@settings(max_examples=5, deadline=None)
@given(note_kw=st.sampled_from(["注：", "参考", "说明", "填表说明", "编制说明"]))
def test_property_note_table_keywords_always_deleted(note_kw):
    """默认关键词命中的注释表总被删除."""
    doc = Document()
    t = doc.add_table(rows=1, cols=1)
    t.rows[0].cells[0].text = f"{note_kw}模板末尾参考内容"
    out, result = fill_and_export(doc, {})
    assert result.note_tables_deleted == 1
    assert len(out.tables) == 0

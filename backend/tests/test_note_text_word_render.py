"""附注正文 HTML → Word：统一渲染与四个出口（第 18 组，design §5.5）。

第 17 组真实写入验收：批量写入（及单章编辑器）存的是 HTML，附注 Word 导出把它当一段纯文字写入 ⇒
``<p>…</p><p>&nbsp;</p>`` 原样出现在 Word 里。本文件锁定：
- 渲染：段落 / 空段 / 加粗倾斜 / 换行 / 列表 / 小标题 / 表格与合并单元格 / 实体 / 未知标签；
- 纯文字在各出口与改造前逐字一致；
- 回填比对投影与交付 docx 提取同口径（逐样本对拍「渲染后再提取」）；
- 四个出口：程序化导出、模板占位填充、交付件刷新、旧版全套填充；交付件回填不把 HTML 章节误判为已修改。
"""

from __future__ import annotations

import io
from types import SimpleNamespace
from unittest.mock import AsyncMock
from uuid import uuid4

import docx
import pytest
from docx.oxml.ns import qn

from app.services import note_text_word_render as render

W_T, W_BR, W_P, W_TBL = qn("w:t"), qn("w:br"), qn("w:p"), qn("w:tbl")

SEPARATED = "<p>本公司以人民币为记账本位币。</p><p>&nbsp;</p><p>外币业务按交易发生日的即期汇率折算。</p>"
TABLE_HTML = ("<p>货币资金明细如下：</p><table><tr><th>项目</th><th>期末余额</th></tr>"
              "<tr><td>库存现金</td><td>1,000.00</td></tr></table><p>无受限资金。</p>")
SAMPLES = [
    SEPARATED,
    TABLE_HTML,
    "<p>普通<strong>加粗</strong><em>斜体</em><b><i>都有</i></b></p>",
    "<p>甲<br>乙</p><p>丙<br/>丁</p>",
    "<ol><li><p>其他应收款按账龄列示</p></li><li>坏账准备</li></ol><ul><li>说明<ul><li>细项</li></ul></li></ul>",
    "<h3>小标题</h3><p>正文</p>",
    "<div>A&amp;B<span class=\"gt-note-placeholder\">C</span><u>D</u></div>",
    "<p>甲</p>\n  <p>乙</p>",
    "前言<p>正文</p>尾声",
    "<table><tr><td colspan=\"2\">合计</td></tr><tr><td>甲</td><td>乙</td></tr></table>",
    "<p></p><p> </p><p>只有这一段有字</p>",
]


def _texts(paragraphs) -> list[str]:
    return [p.text for p in paragraphs]


def _all_text(document) -> str:
    return "\n".join([p.text for p in document.paragraphs] +
                     [c.text for t in document.tables for r in t.rows for c in r.cells])


def _appended(html: str, **kw):
    document = docx.Document()
    render.append_note_text(document, html, **kw)
    return document


class TestDetection:
    @pytest.mark.parametrize("text,expected", [
        ("<p>x</p>", True), ("a<br/>b", True), ("<STRONG>x</STRONG>", True), ("<table><tr><td>1</td></tr></table>", True),
        ("一、公司基本情况\n\n本公司成立于2001年。", False), ("a < b 且 c > d", False), ("<公司名称>", False), ("", False),
        (None, False),
    ])
    def test_is_html_note_text(self, text, expected):
        assert render.is_html_note_text(text) is expected


class TestAppend:
    def test_plain_text_is_one_run_like_before(self):
        """纯文字与改造前逐字一致：去首尾空白后一段一个 run（换行由 python-docx 转为段内换行）。"""
        calls = []
        document = _appended("  一、公司基本情况\n\n本公司成立于2001年。 ",
                             run_style=lambda run, bold: calls.append(bold), paragraph_style=lambda p: calls.append("p"))
        [paragraph] = document.paragraphs
        assert [r.text for r in paragraph.runs] == ["一、公司基本情况\n\n本公司成立于2001年。"]
        assert calls == [False, "p"]

    def test_blank_plain_text_adds_nothing(self):
        assert _appended("  \n ").paragraphs == []

    def test_paragraphs_and_append_separator(self):
        document = _appended(SEPARATED)
        assert _texts(document.paragraphs) == ["本公司以人民币为记账本位币。", "", "外币业务按交易发生日的即期汇率折算。"]
        assert document.paragraphs[1].runs == [], "追加分隔空段不应带 &nbsp; 文字"

    def test_inline_bold_and_italic(self):
        [paragraph] = _appended(SAMPLES[2]).paragraphs
        assert [(r.text, bool(r.bold), bool(r.italic)) for r in paragraph.runs] == [
            ("普通", False, False), ("加粗", True, False), ("斜体", False, True), ("都有", True, True)]

    def test_line_break_stays_inside_paragraph(self):
        paragraphs = _appended(SAMPLES[3]).paragraphs
        assert _texts(paragraphs) == ["甲\n乙", "丙\n丁"]
        assert all(len(p._p.findall(f".//{W_BR}")) == 1 for p in paragraphs)

    def test_whitespace_after_line_break_is_dropped(self):
        """``<br />`` 后的排版换行 / 空格浏览器不显示，Word 里也不应变成行首空格（markdown 转出的 HTML 即此形态）。"""
        assert _texts(_appended("<p>甲<br />\n  乙<br /> </p>").paragraphs) == ["甲\n乙"]

    def test_lists_get_markers_and_li_with_p_is_one_paragraph(self):
        """有序项用附注惯用的「（1）」：交付件回填把「1. 中文」开头的段当章节标题剔除，用「1. 」会让回填误判已修改。"""
        from app.services.deliverable_writeback_service import _is_title_paragraph

        texts = _texts(_appended(SAMPLES[4]).paragraphs)
        assert texts == ["（1）其他应收款按账龄列示", "（2）坏账准备", "• 说明", "  • 细项"]
        assert not [t for t in texts if _is_title_paragraph(t, "八、9")]

    def test_heading_is_bold_but_not_a_heading_style(self):
        """导出带自动目录：正文里的小标题不得用 Word 标题样式（否则进目录）。"""
        heading, body = _appended(SAMPLES[5]).paragraphs
        assert (heading.text, all(r.bold for r in heading.runs), heading.style.name) == ("小标题", True, "Normal")
        assert (body.text, any(r.bold for r in body.runs)) == ("正文", False)

    def test_entities_spans_and_unknown_tags_keep_text(self):
        assert _texts(_appended(SAMPLES[6]).paragraphs) == ["A&BCD"]

    def test_formatting_whitespace_between_blocks_is_not_a_paragraph(self):
        assert _texts(_appended(SAMPLES[7]).paragraphs) == ["甲", "乙"]

    def test_top_level_text_becomes_paragraphs(self):
        assert _texts(_appended(SAMPLES[8]).paragraphs) == ["前言", "正文", "尾声"]

    def test_table_with_header_and_borders(self):
        document = _appended(TABLE_HTML)
        assert _texts(document.paragraphs) == ["货币资金明细如下：", "无受限资金。"]
        body = [el.tag for el in document.element.body if el.tag in (W_P, W_TBL)]
        assert body == [W_P, W_TBL, W_P], "表格应在两段之间（文档顺序）"
        [table] = document.tables
        assert [[c.text for c in row.cells] for row in table.rows] == [["项目", "期末余额"], ["库存现金", "1,000.00"]]
        assert all(r.bold for c in table.rows[0].cells for p in c.paragraphs for r in p.runs)
        assert not any(r.bold for c in table.rows[1].cells for p in c.paragraphs for r in p.runs)
        assert table._tbl.tblPr.find(qn("w:tblBorders")) is not None

    def test_colspan_and_rowspan_merge(self):
        [table] = _appended(SAMPLES[9]).tables
        assert (len(table.rows), len(table.columns)) == (2, 2)
        assert table.cell(0, 0)._tc is table.cell(0, 1)._tc and table.cell(0, 0).text == "合计"
        [table] = _appended("<table><tr><td rowspan=\"2\">类别</td><td>甲</td></tr><tr><td>乙</td></tr></table>").tables
        assert table.cell(0, 0)._tc is table.cell(1, 0)._tc
        assert (table.cell(0, 0).text, table.cell(0, 1).text, table.cell(1, 1).text) == ("类别", "甲", "乙")

    def test_oversized_span_is_clamped(self):
        [table] = _appended("<table><tr><td colspan=\"999\" rowspan=\"999\">x</td></tr></table>").tables
        assert len(table.rows) == 1 and len(table.columns) <= render.MAX_SPAN

    def test_no_markup_leaks_for_any_sample(self):
        for html in SAMPLES:
            text = _all_text(_appended(html))
            assert "<" not in text and "&nbsp;" not in text and "&amp;" not in text, (html, text)


class TestFillPlaceholder:
    def _doc_with_token(self):
        document = docx.Document()
        document.add_paragraph("之前")
        token = document.add_paragraph("{{section:八、1}}")
        token.paragraph_format.left_indent = docx.shared.Pt(21)
        document.add_paragraph("之后")
        return document, token

    def test_blocks_fill_in_place_and_keep_placeholder_format(self):
        document, token = self._doc_with_token()
        render.fill_paragraph_with_note_text(token, TABLE_HTML)
        order = [("tbl" if el.tag == W_TBL else "".join(t.text or "" for t in el.iter(W_T)))
                 for el in document.element.body if el.tag in (W_P, W_TBL)]
        assert order == ["之前", "货币资金明细如下：", "tbl", "无受限资金。", "之后"]
        tail = document.paragraphs[2]
        assert tail.text == "无受限资金。" and tail.paragraph_format.left_indent == docx.shared.Pt(21)

    def test_leading_table_keeps_empty_placeholder_before_it(self):
        document, token = self._doc_with_token()
        render.fill_paragraph_with_note_text(token, SAMPLES[9])
        order = [el.tag for el in document.element.body if el.tag in (W_P, W_TBL)]
        assert order == [W_P, W_P, W_TBL, W_P] and token.text == ""


class TestPlainProjection:
    def test_plain_text_is_identity(self):
        text = "一、公司基本情况\n\n本公司成立于2001年。"
        assert render.note_text_plain(text) == text

    @pytest.mark.parametrize("html", SAMPLES)
    def test_projection_equals_text_extracted_from_rendered_docx(self, html):
        """回填 / 刷新读交付 docx 的口径（``block_text_of``：逐段、跳过表格与空段、段内换行不保留）对拍。"""
        from app.services.section_anchor_utils import block_text_of

        assert render.note_text_plain(html) == block_text_of(render.note_text_elements(html))

    def test_elements_are_detached_and_ordered(self):
        elements = render.note_text_elements(TABLE_HTML)
        assert [el.tag for el in elements] == [W_P, W_TBL, W_P]
        assert all(el.getparent() is None for el in elements)


# ── 四个出口 ──────────────────────────────────────────────────────────────────


def _note(section: str, title: str, text: str) -> SimpleNamespace:
    return SimpleNamespace(note_section=section, section_title=title, text_content=text, table_data=None,
                           guidance_text=None, is_deleted=False, status="draft", is_empty=False, source_template=None)


class TestExporters:
    @pytest.mark.asyncio
    async def test_programmatic_export_renders_html_and_keeps_plain(self):
        from app.services.note_word_exporter import NoteWordExporter

        plain = "本公司会计年度自公历1月1日起至12月31日止。"
        notes = [_note("四、记账本位币", "记账本位币", SEPARATED), _note("四、会计期间", "会计期间", plain)]
        exporter = NoteWordExporter(SimpleNamespace())  # type: ignore[arg-type]
        exporter._load_notes = AsyncMock(return_value=notes)  # type: ignore[method-assign]
        output = await exporter.export(uuid4(), 2025, template_type="soe", report_scope="standalone")
        document = docx.Document(io.BytesIO(output.getvalue()))
        texts = _texts(document.paragraphs)
        assert "本公司以人民币为记账本位币。" in texts and "外币业务按交易发生日的即期汇率折算。" in texts
        assert plain in texts
        assert not [t for t in texts if "<" in t or "&nbsp;" in t], "导出正文含原样 HTML"

    def test_template_placeholder_renders_html_and_keeps_plain(self):
        from app.services.note_word_exporter import NoteWordExporter

        exporter = NoteWordExporter(SimpleNamespace())  # type: ignore[arg-type]
        for text, expected in ((SEPARATED, ["本公司以人民币为记账本位币。", "", "外币业务按交易发生日的即期汇率折算。"]),
                               ("纯文字正文。", ["纯文字正文。"])):
            document = docx.Document()
            token = document.add_paragraph("{{section:八、1}}")
            block = SimpleNamespace(section_code="八、1", elements=[token._element])
            exporter._fill_section_block(document, block, _note("八、1", "货币资金", text))
            assert _texts(document.paragraphs) == expected

    def test_placeholder_sharing_paragraph_gets_text_projection(self):
        """占位符与其它文字同段（四版模板没有，防御夹具与自定义模板）：按文字投影替换，不留标签。"""
        from app.services.note_word_exporter import NoteWordExporter

        document = docx.Document()
        token = document.add_paragraph("说明：{{section:八、1}}")
        block = SimpleNamespace(section_code="八、1", elements=[token._element])
        NoteWordExporter(SimpleNamespace())._fill_section_block(document, block, _note("八、1", "货币资金", SEPARATED))  # type: ignore[arg-type]
        assert _texts(document.paragraphs) == ["说明：本公司以人民币为记账本位币。\n外币业务按交易发生日的即期汇率折算。"]

    def test_refresh_replaces_block_with_rendered_html(self):
        from app.services.deliverable_refresh_service import DeliverableRefreshService
        from app.services.section_anchor_utils import block_text_hash
        from app.services.word_doc_utils import scan_section_blocks

        for text, expected in ((TABLE_HTML, ["货币资金明细如下：", "无受限资金。"]), ("第一行\n第二行", ["第一行", "第二行"])):
            document = docx.Document()
            for line in ("##SECTION:八、1##", "旧正文", "##/SECTION:八、1##"):
                document.add_paragraph(line)
            [block] = scan_section_blocks(document)
            new_hash = DeliverableRefreshService(SimpleNamespace())._replace_block_content(block, text, "marker")
            inner = _texts(document.paragraphs)[1:-1]
            assert inner == expected, inner
            [refreshed] = scan_section_blocks(document)
            assert new_hash == block_text_hash(refreshed.elements)
            assert len(document.tables) == (1 if text == TABLE_HTML else 0)


class TestWritebackComparison:
    async def _diff(self, block_texts_or_html: list[str], db_text: str, *, fill_html: str | None = None) -> dict:
        from app.services.deliverable_writeback_service import DeliverableWritebackService

        document = docx.Document()
        document.add_paragraph("##SECTION:四、记账本位币##")
        if fill_html is not None:
            render.fill_paragraph_with_note_text(document.add_paragraph("{{section:四、记账本位币}}"), fill_html)
        for text in block_texts_or_html:
            document.add_paragraph(text)
        document.add_paragraph("##/SECTION:四、记账本位币##")
        buf = io.BytesIO()
        document.save(buf)
        service = DeliverableWritebackService(AsyncMock())
        extracted = await service._extract_sections_from_docx(buf.getvalue())
        return service._compute_section_diff(extracted, {"四、记账本位币": db_text})

    @pytest.mark.asyncio
    async def test_rendered_html_chapter_is_not_a_change(self):
        html = "<ol><li>本公司以人民币为记账本位币。</li></ol>" + SEPARATED + TABLE_HTML
        assert await self._diff([], html, fill_html=html) == {}

    @pytest.mark.asyncio
    async def test_legacy_raw_html_deliverable_is_not_a_change(self):
        """修复前生成的交付件：正文段里是原样 HTML 字符串，与库里同一 HTML 比较不得判为已修改。"""
        assert await self._diff([SEPARATED], SEPARATED) == {}

    @pytest.mark.asyncio
    async def test_real_edit_is_still_detected(self):
        diff = await self._diff(["本公司以美元为记账本位币。"], SEPARATED)
        assert list(diff) == ["四、记账本位币"]

    def test_plain_text_comparison_unchanged(self):
        from app.services.deliverable_writeback_service import DeliverableWritebackService

        service = DeliverableWritebackService(AsyncMock())
        assert service._compute_section_diff({"八、1": "甲  乙"}, {"八、1": "甲 乙"}) == {}
        assert list(service._compute_section_diff({"八、1": "甲"}, {"八、1": "乙"})) == ["八、1"]


# ── 旧版全套填充（routers/word_export.py full-package 仍在用）：复用其测试的内存库夹具 ──────────────

from tests import test_phase13_word_export as _phase13  # noqa: E402

test_db, test_user, test_project = _phase13.test_db, _phase13.test_user, _phase13.test_project


class TestLegacyFiller:
    @pytest.mark.asyncio
    async def test_fill_disclosure_notes_renders_html_and_keeps_plain(self, test_db, test_project, test_user):
        import shutil

        from app.models.report_models import DisclosureNote
        from app.services.word_template_filler import STORAGE_ROOT, WordTemplateFiller

        plain = "本公司会计年度自公历1月1日起至12月31日止。"
        test_db.add_all([
            DisclosureNote(project_id=test_project.id, year=2024, note_section="四、记账本位币", section_title="记账本位币",
                           text_content=SEPARATED, sort_order=1),
            DisclosureNote(project_id=test_project.id, year=2024, note_section="四、会计期间", section_title="会计期间",
                           text_content=plain, sort_order=2),
        ])
        await test_db.flush()
        try:
            path = await WordTemplateFiller(test_db).fill_disclosure_notes(test_db, test_project.id, 2024, test_user.id)
            texts = _texts(docx.Document(str(path)).paragraphs)
        finally:
            shutil.rmtree(STORAGE_ROOT / "projects" / str(test_project.id), ignore_errors=True)
        assert "本公司以人民币为记账本位币。" in texts and "外币业务按交易发生日的即期汇率折算。" in texts
        assert plain in texts
        assert not [t for t in texts if "<p>" in t or "&nbsp;" in t], "旧版填充正文含原样 HTML"

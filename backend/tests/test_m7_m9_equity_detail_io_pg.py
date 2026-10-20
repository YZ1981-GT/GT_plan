"""M7-2 / M9-2 导入导出与 M9 公式校验 —— 真实 PostgreSQL + 真实应用路由 + 真实 Wp_Bound_Gate。

SQLite 用例（``test_m7_m9_equity_detail_io.py``）证明 service 逻辑与前端契约；这里证明只有 PG 才能证明的：
``(wp_id, item_id)`` 唯一约束上的 ON CONFLICT、uuid 列绑定字符串参数、``project_id`` 非空与
``updated_by`` 外键，以及生产注册的 router 级 dedicated_wp_gate（admin 放行、项目外用户 404）。

每个用例在一个最终回滚的事务里建数据；会话用 ``create_savepoint`` 吸收路由里的 commit，不污染开发库。
"""

from __future__ import annotations

import io
import json

import pytest
import pytest_asyncio
import sqlalchemy as sa
from openpyxl import load_workbook
from sqlalchemy.ext.asyncio import AsyncSession, create_async_engine

from app.core.config import settings
from app.services import m7_special_reserve_service as m7
from app.services import m9_other_comprehensive_income_service as m9
from app.services import m_detail_sheet_io as core
from tests.procedure_delegation_visibility._factories import mk_project, mk_user, mk_working_paper, mk_wp_index

IS_PG = settings.DATABASE_URL.startswith("postgresql")
pytestmark = [pytest.mark.pg_only, pytest.mark.skipif(not IS_PG, reason="need PostgreSQL")]

_XLSX = "application/vnd.openxmlformats-officedocument.spreadsheetml.sheet"


@pytest_asyncio.fixture
async def pg_session():
    engine = create_async_engine(settings.DATABASE_URL, pool_pre_ping=True)
    conn = await engine.connect()
    trans = await conn.begin()
    session = AsyncSession(bind=conn, expire_on_commit=False, join_transaction_mode="create_savepoint")
    try:
        yield session
    finally:
        await session.close()
        await trans.rollback()
        await conn.close()
        await engine.dispose()


async def _workpaper(session, wp_code: str):
    project = await mk_project(session)
    index = await mk_wp_index(session, project.id, wp_code=wp_code, audit_cycle="M")
    paper = await mk_working_paper(session, project.id, index.id)
    return project, paper


def _filled(spec, rows) -> bytes:
    wb = load_workbook(io.BytesIO(core.export_template([spec]).getvalue()))
    ws = wb[spec.sheet_title]
    for row in rows:
        ws.append(list(row))
    buffer = io.BytesIO()
    wb.save(buffer)
    return buffer.getvalue()


def _payload(resp):
    body = resp.json()
    return body.get("data", body) if isinstance(body, dict) else body


async def _rows(session, wp_id) -> dict:
    result = await session.execute(
        sa.text(
            "SELECT item_id, conclusion, remark, project_id, updated_by "
            "FROM checklist_responses WHERE wp_id = :w"
        ),
        {"w": wp_id},
    )
    return {row.item_id: row for row in result.fetchall()}


async def test_m7_import_and_export_through_real_routes(pg_session):
    from app.main import app
    from tests._test_auth_helper import FakeAuthUser, override_auth

    project, paper = await _workpaper(pg_session, "M7")
    admin = await mk_user(pg_session, role="admin")
    base = f"/api/m7-special-reserve/{paper.id}"
    rows = [("安全生产费", 1000, 500, 200, 100, 10, 0, 5, 0, 3, 0, 800, 400, 200, 1000, "备注"), ("维简费", 300)]
    async with override_auth(app, db_session=pg_session, user=FakeAuthUser(user_id=admin.id)) as client:
        first = await client.post(
            f"{base}/import-data",
            files={"file": ("m7.xlsx", _filled(m7.M7_DETAIL_SPEC, rows), _XLSX)},
            data={"sheet": "M7-2"},  # 与前端同形：sheet 在表单字段里
        )
        assert first.status_code == 200, first.text
        assert _payload(first)["imported_count"] == 2
        # 再导入 1 行：ON CONFLICT 走真实唯一约束，超出本次行数的逐行键清空
        second = await client.post(
            f"{base}/import-data", files={"file": ("m7.xlsx", _filled(m7.M7_DETAIL_SPEC, rows[1:]), _XLSX)}
        )
        assert second.status_code == 200, second.text
        exported = await client.get(f"{base}/export-data", params={"sheet": "M7-2"})
        assert exported.status_code == 200, exported.text
        unknown = await client.get(f"{base}/export-template", params={"sheet": "M7-4"})
        assert unknown.status_code == 400, unknown.text
        accrual = await client.post(f"{base}/accrual-test", json={"method": "output", "tiers": [{"output": 10, "rate": 3}]})
        assert accrual.status_code == 200 and _payload(accrual)["estimated"] == 30

    stored = await _rows(pg_session, paper.id)
    assert {row.project_id for row in stored.values()} == {project.id}
    assert {row.updated_by for row in stored.values()} == {admin.id}
    assert stored["M7-2-full-data"].conclusion is None
    assert [r["itemName"] for r in json.loads(stored["M7-2-full-data"].remark)] == ["维简费"]
    assert stored["M7-2-row-1-data"].remark is None and stored["M7-2-row-1-auditedEnd"].remark is None
    assert stored["M7-2-row-0-auditedEnd"].remark == "300" and stored["M7-2-auditedEndTotal"].remark == "300"
    sheet = load_workbook(io.BytesIO(exported.content), data_only=True)["明细表M7-2"]
    assert [row[0] for row in sheet.iter_rows(min_row=2, values_only=True)] == ["维简费", "合计"]


async def test_m9_import_and_formula_check_through_real_routes(pg_session):
    from app.main import app
    from tests._test_auth_helper import FakeAuthUser, override_auth

    project, paper = await _workpaper(pg_session, "M9")
    admin = await mk_user(pg_session, role="admin")
    for key, unadjusted in (("m9-adj-a", 120), ("m9-adj-b", 200)):
        await pg_session.execute(
            sa.text("INSERT INTO checklist_responses (project_id, wp_id, item_id, remark) VALUES (:p, :w, :i, :r)"),
            {"p": project.id, "w": paper.id, "i": f"M9-1-{key}-data", "r": json.dumps(
                {"key": key, "itemName": key, "beginning": 100, "creditAmount": 30, "debitAmount": 10,
                 "unadjusted": unadjusted})},
        )
    base = f"/api/m9-other-comprehensive-income/{paper.id}"
    detail = [("其他权益工具投资公允价值变动", "不可重分类", 1000, 400, 100, 400, 100)]
    async with override_auth(app, db_session=pg_session, user=FakeAuthUser(user_id=admin.id)) as client:
        imported = await client.post(
            f"{base}/import-data", files={"file": ("m9.xlsx", _filled(m9.M9_DETAIL_SPEC, detail), _XLSX)}
        )
        assert imported.status_code == 200, imported.text
        checked = await client.get(f"{base}/validate-formulas")
        assert checked.status_code == 200, checked.text

    result = _payload(checked)
    assert (result["total_rows"], result["valid_rows"], result["invalid_rows"]) == (2, 1, 1)
    assert {r["row_key"]: r["is_valid"] for r in result["results"]} == {"m9-adj-a": True, "m9-adj-b": False}
    stored = await _rows(pg_session, paper.id)
    assert stored["M9-2-detail-rows"].remark is None
    assert json.loads(stored["M9-2-detail-rows"].conclusion)[0]["ociCategory"] == "nonReclass"
    assert stored["M9-2-auditedEndTotal"].remark == "1300"


async def test_gate_hides_the_routes_from_users_outside_the_project(pg_session):
    from app.main import app
    from tests._test_auth_helper import FakeAuthUser, override_auth

    _project, paper = await _workpaper(pg_session, "M7")
    outsider = await mk_user(pg_session, role="auditor")
    user = FakeAuthUser(user_id=outsider.id, username=outsider.username, role="auditor")
    content = _filled(m7.M7_DETAIL_SPEC, [("安全生产费", 1)])
    async with override_auth(app, db_session=pg_session, user=user) as client:
        read = await client.get(f"/api/m7-special-reserve/{paper.id}/export-data")
        write = await client.post(
            f"/api/m7-special-reserve/{paper.id}/import-data", files={"file": ("m7.xlsx", content, _XLSX)}
        )
        check = await client.get(f"/api/m9-other-comprehensive-income/{paper.id}/validate-formulas")
    assert (read.status_code, write.status_code, check.status_code) == (404, 404, 404), (read.text, write.text)
    assert await _rows(pg_session, paper.id) == {}

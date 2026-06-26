"""Property-Based Test: word-template 模板文件可解析 (Property 2)

对于任意注册为 "word-template" 的 wp_code，
验证 backend/wp_templates/A/ 中存在对应 .docx 文件，
且 python-docx 可正常打开，文档包含至少一个段落或表格（非空）。

Feature: a-cycle-docx-online, Property 2: word-template 模板文件可解析

**Validates: Requirements 1.4, 1.5**
"""
from __future__ import annotations

from pathlib import Path

from docx import Document
from hypothesis import given, settings
from hypothesis import strategies as st

# ──────────────────────────────────────────────────────────────────────
# 25 个 word-template wp_code（与 wp_code_overrides.json 保持一致）
# ──────────────────────────────────────────────────────────────────────

WORD_TEMPLATE_CODES: list[str] = [
    "A8-1", "A8-2", "A9-1", "A9-2", "A10-1", "A11-1", "A12-1",
    "A16-1", "A16-2", "A16-3", "A16-4", "A16-5", "A16-6", "A16-7",
    "A17-2-1", "A17-3", "A17-3-1", "A17-4", "A17-6",
    "A18-1", "A26-1", "A26-2", "A26-3", "A26-4", "A27-1",
]

# ──────────────────────────────────────────────────────────────────────
# 模板目录
# ──────────────────────────────────────────────────────────────────────

_TEMPLATE_DIR = Path(__file__).resolve().parent.parent / "wp_templates" / "A"


# ──────────────────────────────────────────────────────────────────────
# Helper：根据 wp_code 查找对应的 .docx 模板文件
# ──────────────────────────────────────────────────────────────────────


def _find_template_file(wp_code: str) -> Path | None:
    """在 wp_templates/A/ 目录中查找以 wp_code 开头的 .docx 文件。

    文件命名规则：{wp_code} + 描述文字 + .docx
    例如: "A8-1 管理层对审计报告日后公布其他信息的书面声明201707.docx"
    """
    for f in _TEMPLATE_DIR.iterdir():
        if not f.suffix.lower() == ".docx":
            continue
        # 文件名以 wp_code 开头（后跟空格或非字母数字字符）
        name = f.name
        if name.startswith(wp_code) and (
            len(name) == len(wp_code) + 5  # 恰好 "{wp_code}.docx"
            or name[len(wp_code)] in (" ", "\u3000", "-", "_", ".")
        ):
            return f
    return None


# ──────────────────────────────────────────────────────────────────────
# Property-Based Test
# ──────────────────────────────────────────────────────────────────────


@given(wp_code=st.sampled_from(WORD_TEMPLATE_CODES))
@settings(max_examples=5, deadline=None)
def test_template_file_exists_and_parseable(wp_code: str):
    """Property 2: 任意 word-template wp_code 对应模板文件存在且 python-docx 可打开

    Feature: a-cycle-docx-online, Property 2: word-template 模板文件可解析

    **Validates: Requirements 1.4, 1.5**
    """
    # 1. 验证模板目录存在
    assert _TEMPLATE_DIR.exists(), f"模板目录不存在: {_TEMPLATE_DIR}"

    # 2. 验证对应 .docx 文件存在
    template_path = _find_template_file(wp_code)
    assert template_path is not None, (
        f"wp_code {wp_code!r} 在 {_TEMPLATE_DIR} 中找不到对应 .docx 文件"
    )
    assert template_path.exists(), (
        f"模板文件路径已定位但文件不存在: {template_path}"
    )

    # 3. 验证 python-docx 可打开（不抛异常）
    doc = Document(str(template_path))

    # 4. 验证文档非空（至少有一个段落或表格）
    has_paragraphs = len(doc.paragraphs) > 0
    has_tables = len(doc.tables) > 0
    assert has_paragraphs or has_tables, (
        f"模板文件 {template_path.name} 为空文档（无段落无表格）"
    )

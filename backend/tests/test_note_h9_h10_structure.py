"""附注 H9 租赁负债 + H10 资产处置损益章节结构守卫。

章节：H10 上市 `三、资产处置收益（损` / 国企 `八、75`；H9 上市 `五、47` / 国企 `八、52`。

锁定 `fix_note_h9_h10_structure.py` 的对齐结果，并固化两个 P0 的修复：

1. **上市两张表原本同名 `项  目`** → `sub_table_data` 以表名为键，同名互相覆盖丢整张表
   → 正名为 `资产处置收益（损失以“-”填列）` / `试运行销售损益`，守卫断言 `项  目` 不复活。
2. **上市 `section_number` 是 md 截断值 `三、资产处置收益（损`（10 字符）** —— 这是既有真源
   形态，前端常量必须对齐它（守卫从前端源码抽常量交叉比对，防有人"修正"成完整名）。

spec: h9-h10-remaining-disclosure-alignment (Task 5)
"""
from __future__ import annotations

import importlib.util
import json
import re
from pathlib import Path

import openpyxl

_ROOT = Path(__file__).resolve().parent.parent
_FIX = _ROOT / "scripts" / "fix" / "fix_note_h9_h10_structure.py"
_H9_XLSX = _ROOT / "wp_templates" / "H" / "H9 租赁负债.xlsx"
_H10_XLSX = _ROOT / "wp_templates" / "H" / "H10 资产处置损益.xlsx"
_FE = (
    _ROOT.parent / "audit-platform" / "frontend" / "src" / "components"
    / "workpaper" / "composables"
)
_FE_H10_MAP = _FE / "h10NoteSectionMap.ts"
_FE_H10_PAYLOAD = _FE / "h10DisclosureSyncPayload.ts"
_FE_H9_PAYLOAD = _FE / "h9DisclosureSyncPayload.ts"

_SHEET_LISTED = "附注披露信息（上市公司）"
_SHEET_SOE = "附注披露信息（国企）"
_LEAKED = "项  目"


def _load_fix():
    spec = importlib.util.spec_from_file_location("_fix_h9h10_struct", _FIX)
    mod = importlib.util.module_from_spec(spec)
    spec.loader.exec_module(mod)
    return mod


FIX = _load_fix()


def _norm(v: object) -> str:
    return re.sub(r"\s+", "", str(v or ""))


def _strip_ts_comments(src: str) -> str:
    """去掉 TS 注释再做源码断言。

    🔴 必须做：注释里会写「不再需要 `项  目__trial`」这类**反例说明**，
    不剥离会被数成真实代码（本守卫首轮即因此误报两条）。
    配套 `test_reverse_self_check_comment_stripper` 防剥离过头导致断言空转。
    """
    src = re.sub(r"/\*[\s\S]*?\*/", "", src)
    return re.sub(r"(^|[^:])//.*$", r"\1", src, flags=re.M)


def _section(key: str) -> dict:
    path, section_number, _plan, _ts = FIX._TARGETS[key]
    doc = json.loads(Path(path).read_text(encoding="utf-8"))
    sec = next(
        (s for s in doc["sections"] if str(s.get("section_number")) == section_number), None
    )
    assert sec is not None, f"{key} 章节 {section_number!r} 不存在"
    return sec


def _tables(key: str) -> dict[str, dict]:
    return {str(t.get("name")): t for t in (_section(key).get("tables") or [])}


def _src_col1(xlsx: Path, sheet: str, first: int, last: int) -> list[str]:
    wb = openpyxl.load_workbook(xlsx, data_only=True)
    assert sheet in wb.sheetnames, wb.sheetnames
    ws = wb[sheet]
    return [_norm(ws.cell(r, 1).value) for r in range(first, last + 1)]


# ───────────────────── Property 8：幂等与零欠账 ─────────────────────


def test_check_passes_all_sections():
    for key in FIX._TARGETS:
        _changes, warnings, errs = FIX._runner(key, dry_run=True, check=True)
        assert not errs, f"{key} 结构欠账：{errs}"
        assert not warnings, f"{key} 告警：{warnings}"


def test_templates_parse():
    for path in {t[0] for t in FIX._TARGETS.values()}:
        json.loads(Path(path).read_text(encoding="utf-8"))


# ───────────── Property 1/2：章节号与表名（含 `项  目` 不复活）─────────────


def test_section_numbers_are_verbatim():
    """listed 侧是 md 截断值，禁"修正"（改了会让同步落空并新建垃圾章节）。"""
    assert FIX.H10_LISTED_SECTION == "三、资产处置收益（损"
    assert len(FIX.H10_LISTED_SECTION) == 10
    assert FIX.H10_SOE_SECTION == "八、75"
    assert FIX.H9_LISTED_SECTION == "五、47"
    assert FIX.H9_SOE_SECTION == "八、52"


def test_h10_listed_tables_renamed():
    names = list(_tables("h10_listed"))
    assert names == [FIX.H10_LISTED_MAIN, FIX.H10_LISTED_TRIAL], names
    assert _LEAKED not in names, "表头首格泄漏名复活 → 两张表又会同名互相覆盖"
    assert len(set(names)) == len(names), "表名重复"


def test_no_leaked_table_name_anywhere():
    for key in FIX._TARGETS:
        assert _LEAKED not in _tables(key), f"{key} 残留 `{_LEAKED}`"


def test_h10_soe_and_h9_table_names():
    assert list(_tables("h10_soe")) == [FIX.H10_SOE_MAIN]
    assert list(_tables("h9_listed")) == [FIX.H9_MAIN]
    assert list(_tables("h9_soe")) == [FIX.H9_MAIN]


def test_frontend_constants_match_template():
    """前端常量 ↔ 模板/源 xlsx 逐字（章节号 / sheet 名 / 表名 / 旧名清单）。"""
    src = _FE_H10_MAP.read_text(encoding="utf-8")
    # 🔴 章节号字面量必须**内联**在 `H10_NOTE_SECTION` 对象里：
    # `gen_note_wp_sync_registry.py` 只认 `listed: '…'` 形态的字符串字面量，
    # 写成标识符引用会让 H10 整条从 registry 消失（本 spec 实测踩过）。
    assert f"listed: '{FIX.H10_LISTED_SECTION}'" in src, "上市章节号漂移或改成了标识符引用"
    assert "H10_LISTED_NOTE_SECTION: string = H10_NOTE_SECTION.listed" in src, (
        "上市章节号应从 H10_NOTE_SECTION 派生（单一真源），不得另写一份字面量"
    )
    assert "H10_SOE_NOTE_SECTION = '八、75'" in src
    # 源 xlsx tab 名：H10 国企是「国企」，不是「国有企业」（H7 才是国有企业）
    assert f"soe: '{_SHEET_SOE}'" in src, "国企 sheet 名未对齐源 xlsx tab 名"
    assert "附注披露信息（国有企业）" not in src, "国企 sheet 名仍为旧的「国有企业」"
    assert f"listed: '{FIX.H10_LISTED_MAIN}'" in src, "上市主表名漂移"
    assert f"soe: '{FIX.H10_SOE_MAIN}'" in src, "国企主表名漂移"
    assert f"H10_TRIAL_SUBTABLE = '{FIX.H10_LISTED_TRIAL}'" in src, "试运行表名漂移"
    assert f"H10_LEGACY_OBSOLETE_TABLES: readonly string[] = ['{_LEAKED}']" in src, (
        "旧名未登记 → 附注会残留孤儿表"
    )


def test_h10_note_section_object_body_has_no_comment():
    """🔴 `H10_NOTE_SECTION` 对象体内不得写注释。

    `gen_note_wp_sync_registry.py` 用 `listed\\s*:\\s*'…'` 从对象体抽章节号，
    注释里若出现 `listed: '…'` 字样会被**优先**抓到（实测抓到 `…` 后
    `_is_section_code` 判否 → 整条 H10 从 registry 消失）。
    """
    src = _FE_H10_MAP.read_text(encoding="utf-8")
    m = re.search(r"H10_NOTE_SECTION\s*=\s*\{(?P<body>[^}]*)\}", src, re.S)
    assert m, "未定位到 H10_NOTE_SECTION 对象"
    body = m.group("body")
    assert "//" not in body and "/*" not in body, f"对象体内含注释：{body!r}"
    assert body.count("listed") == 1, f"对象体内 `listed` 出现多次：{body!r}"


def test_registry_reflects_h10_constants():
    """registry 是「同步就绪度看板」的定位真源，须与前端常量一致（改常量后需重生成）。"""
    reg = json.loads(
        (_ROOT / "data" / "note_workpaper_sync_registry.json").read_text(encoding="utf-8")
    )
    hit = next((e for e in reg["entries"] if e["wp_code"] == "H10"), None)
    assert hit is not None, "registry 缺 H10 条目 —— 生成器未抽到章节号（多为字面量被改成标识符引用）"
    assert hit["listed"] == FIX.H10_LISTED_SECTION, hit
    assert hit["soe"] == FIX.H10_SOE_SECTION, hit
    assert hit["sheet_soe"] == _SHEET_SOE, hit
    assert hit["sheet_listed"] == _SHEET_LISTED, hit
    h9 = next((e for e in reg["entries"] if e["wp_code"] == "H9"), None)
    assert h9 and h9["listed"] == "五、47" and h9["soe"] == "八、52", h9


def test_frontend_payload_dropped_bypass_keys():
    """`项  目__trial` / `_trial_detail` 是孤儿绕过键（模板里无此表名）→ 必须删净。"""
    src = _strip_ts_comments(_FE_H10_PAYLOAD.read_text(encoding="utf-8"))
    assert "__trial" not in src, "残留 `__trial` 后缀绕过键"
    assert "_trial_detail" not in src, "残留 `_trial_detail` 双写"
    assert "_removed_table_keys" in src, "改名后须发 _removed_table_keys"
    assert "H10_NOTE_TEXT_TITLES" in src and "buildH10NoteTexts" in src


# ───────── Property 3：行集与源 xlsx 一致且无假行/占位行 ─────────


def test_h10_listed_main_rows():
    rows = _tables("h10_listed")[FIX.H10_LISTED_MAIN].get("rows") or []
    labels = [_norm(r.get("label")) for r in rows]
    assert labels == [_norm(x) for x in FIX.H10_LISTED_ROW_LABELS + ["合计"]], labels
    assert len(rows) == 11, len(rows)
    assert rows[-1].get("is_total") is True


def test_h10_soe_main_rows_match_source_xlsx():
    """国企 R9~R18 十行 + R20 `合  计`（R19 为空行）。"""
    rows = _tables("h10_soe")[FIX.H10_SOE_MAIN].get("rows") or []
    labels = [_norm(r.get("label")) for r in rows]
    src = [x for x in _src_col1(_H10_XLSX, _SHEET_SOE, 9, 20) if x]
    assert src == [_norm(x) for x in FIX.H10_SOE_ROW_LABELS] + ["合计"], src
    assert labels == src, f"模板={labels}\n源={src}"


def test_h10_listed_main_rows_match_source_xlsx_modulo_quote_width():
    """上市 R9~R18 十行；源模板三行用半角 `"-"`，模板统一全角 —— 去掉引号后必须逐字相等。"""
    src = _src_col1(_H10_XLSX, _SHEET_LISTED, 9, 18)
    strip_q = lambda s: re.sub(r"[\"'“”]", "", s)  # noqa: E731
    assert [strip_q(x) for x in src] == [
        strip_q(_norm(x)) for x in FIX.H10_LISTED_ROW_LABELS
    ], src
    # 模板侧引号宽度必须**统一为全角**（同表内混用会让行标签匹配随机失败）
    rows = _tables("h10_listed")[FIX.H10_LISTED_MAIN].get("rows") or []
    for r in rows:
        assert '"' not in str(r.get("label")), f"残留半角引号：{r.get('label')}"


def test_h10_listed_trial_rows_match_source_xlsx():
    rows = _tables("h10_listed")[FIX.H10_LISTED_TRIAL].get("rows") or []
    labels = [_norm(r.get("label")) for r in rows]
    src = [x for x in _src_col1(_H10_XLSX, _SHEET_LISTED, 29, 31) if x]
    assert src == ["固定资产试运行销售", "研发样品销售", "合计"], src
    assert labels == src, labels


def test_no_placeholder_or_fake_rows():
    for key in FIX._TARGETS:
        for name, tbl in _tables(key).items():
            for r in tbl.get("rows") or []:
                assert r.get("row_type") != "header_label", f"{key}/{name} 残留 header_label 假行"
                assert "可无限量添加行" not in _norm(r.get("label")), (
                    f"{key}/{name} 残留占位说明行"
                )


def test_h9_rows_untouched():
    """H9 结构本就对：上市 3 行、国企 5 行含 `……` 可扩扣减行（须保留）。"""
    listed = _tables("h9_listed")[FIX.H9_MAIN].get("rows") or []
    assert [_norm(r.get("label")) for r in listed] == [
        "小计", "减：一年内到期的租赁负债", "合计",
    ]
    soe = _tables("h9_soe")[FIX.H9_MAIN].get("rows") or []
    assert [_norm(r.get("label")) for r in soe] == [
        "租赁付款额", "减：未确认的融资费用", "重分类至一年内到期的非流动负债",
        "……", "租赁负债净额",
    ]
    src = [x for x in _src_col1(_H9_XLSX, _SHEET_SOE, 8, 12)]
    assert src == [_norm(r.get("label")) for r in soe], src


# ───────── Property 4：两级表头正确、单级表 flat ─────────


def test_h10_listed_trial_two_level_columns():
    tbl = _tables("h10_listed")[FIX.H10_LISTED_TRIAL]
    cols = tbl.get("columns") or []
    assert [c.get("key") for c in cols] == [
        "label", "current_income", "current_cost", "prior_income", "prior_cost",
    ], cols
    assert [c.get("label") for c in cols] == tbl["headers"]
    assert [c.get("label") for c in cols] == ["项目", "收入", "成本", "收入", "成本"]
    assert cols[0].get("is_label") is True
    assert not cols[0].get("group"), "标签列 rowspan=2 不得带 group"
    assert not any(c.get("flat") for c in cols), "两级表头不得标 flat"
    assert tbl.get("_column_groups") == [
        {"group": "本期发生额", "start": 1, "span": 2},
        {"group": "上期发生额", "start": 3, "span": 2},
    ], tbl.get("_column_groups")
    # 与源 xlsx R27/R28 两级表头交叉比对
    wb = openpyxl.load_workbook(_H10_XLSX, data_only=True)
    ws = wb[_SHEET_LISTED]
    assert [_norm(ws.cell(27, c).value) for c in (1, 2, 4)] == ["项目", "本期发生额", "上期发生额"]
    assert [_norm(ws.cell(28, c).value) for c in (2, 3, 4, 5)] == ["收入", "成本", "收入", "成本"]


def test_flat_single_level_tables():
    cases = [
        ("h10_listed", FIX.H10_LISTED_MAIN, ["label", "current_amount", "prior_amount"]),
        ("h10_soe", FIX.H10_SOE_MAIN,
         ["label", "current_amount", "prior_amount", "non_recurring_amount"]),
        ("h9_listed", FIX.H9_MAIN, ["label", "end_balance", "prior_balance"]),
        ("h9_soe", FIX.H9_MAIN, ["label", "end_balance", "begin_balance"]),
    ]
    for key, name, keys in cases:
        tbl = _tables(key)[name]
        cols = tbl.get("columns") or []
        assert [c.get("key") for c in cols] == keys, (key, cols)
        assert [c.get("label") for c in cols] == tbl["headers"], key
        assert cols[0].get("is_label") and cols[0].get("flat"), f"{key} 首列须 is_label + flat"
        assert all(c.get("format") == "amount" for c in cols[1:]), key
        assert not any(c.get("group") for c in cols), f"{key} 单级表头不得声明 group"
        assert not tbl.get("_column_groups"), f"{key} 单级表头不得残留 _column_groups"


def test_source_headers_cross_check():
    wb10 = openpyxl.load_workbook(_H10_XLSX, data_only=True)
    assert [_norm(wb10[_SHEET_LISTED].cell(8, c).value) for c in range(1, 4)] == [
        "项目", "本期发生额", "上期发生额",
    ]
    assert [_norm(wb10[_SHEET_SOE].cell(8, c).value) for c in range(1, 5)] == [
        "项目", "本期发生额", "上期发生额", "计入当期非经常性损益的金额",
    ]
    wb9 = openpyxl.load_workbook(_H9_XLSX, data_only=True)
    assert [_norm(wb9[_SHEET_LISTED].cell(7, c).value) for c in range(1, 4)] == [
        "项目", "期末余额", "上年年末余额",
    ]
    assert [_norm(wb9[_SHEET_SOE].cell(7, c).value) for c in range(1, 4)] == [
        "项目", "期末余额", "期初余额",
    ]


# ───────── Property 7：guidance 非空 + text_sections 清洁 ─────────


def test_guidance_present():
    for key in FIX._TARGETS:
        for name, tbl in _tables(key).items():
            assert str(tbl.get("guidance") or "").strip(), f"{key}/{name} 缺 guidance"


def test_h10_soe_text_sections_no_markdown_marks():
    paras = _section("h10_soe").get("text_sections") or []
    assert paras, "国企说明段被清空"
    for p in paras:
        assert "**" not in str(p), f"残留 markdown 残迹：{str(p)[:40]}"
    # 源模板 R25 第 3 条原本整段缺失
    assert any("其他业务收入/成本" in str(p) for p in paras), "缺源模板第 3 条"


def test_h10_listed_text_sections():
    paras = [str(p) for p in (_section("h10_listed").get("text_sections") or [])]
    assert paras[0] == "【提示：", paras[0]
    assert any("4.使用权资产、油气资产处置损益" in p for p in paras), "缺源模板第 4 条"
    # 裸表名会被当披露正文渲染 → 必须带 `#### `
    assert f"#### {FIX.H10_LISTED_TRIAL}" in paras
    assert FIX.H10_LISTED_TRIAL not in paras, "裸表名段落未加 #### 前缀"


def test_h9_note_text_titles_registered():
    src = _strip_ts_comments(_FE_H9_PAYLOAD.read_text(encoding="utf-8"))
    assert "H9_NOTE_TEXT_TITLES" in src and "buildH9NoteTexts" in src
    for section in ("listed-interest", "soe-guidance"):
        assert f"'{section}':" in src, f"缺 {section} 中文标题"
    assert src.count("flat: true") == 2, "两版列首列须各标一次 flat（上市 + 国企）"


# ───────────────────── 反向自检（防守卫空转） ─────────────────────


def test_reverse_self_check_validate_section():
    from _note_structure_kit import validate_section

    broken = {
        "tables": [
            {"name": FIX.H9_MAIN, "headers": ["项目"], "rows": [], "guidance": "x"}
        ]
    }
    errs = validate_section(broken, [FIX.H9_MAIN])
    assert any("缺 columns" in e for e in errs), errs


def test_reverse_self_check_duplicate_name_detected():
    """同名表必须被 validate_section 抓出（这正是 H10 上市原本的 P0）。"""
    from _note_structure_kit import validate_section

    dup = {
        "tables": [
            {"name": _LEAKED, "headers": ["项目"], "columns": [], "rows": [], "guidance": ""},
            {"name": _LEAKED, "headers": ["项目"], "columns": [], "rows": [], "guidance": ""},
        ]
    }
    errs = validate_section(dup, [_LEAKED])
    assert any("表名重复" in e for e in errs), errs


def test_reverse_self_check_normalizer():
    assert _norm("合  计") == "合计"
    assert _norm("    1.期初余额") == "1.期初余额"
    assert _norm(None) == ""


def test_reverse_self_check_comment_stripper():
    """`_strip_ts_comments` 既要剥掉注释、又不能把代码一起吃掉（否则断言空转）。"""
    sample = (
        "/* 不再需要 `项  目__trial` */\n"
        "const a = 1 // _trial_detail 已删\n"
        "const url = 'https://x/y'\n"
        "const cols = [{ flat: true }]\n"
    )
    out = _strip_ts_comments(sample)
    assert "__trial" not in out and "_trial_detail" not in out
    assert "const a = 1" in out and "flat: true" in out
    assert "https://x/y" in out, "不得把 URL 里的 // 当注释切掉"
    # 原始源码里确实含被禁字样 → 证明上面两条断言不是空转
    raw = _FE_H10_PAYLOAD.read_text(encoding="utf-8")
    assert "__trial" in raw and "_trial_detail" in raw, (
        "载荷源码注释里已不再提这两个绕过键 → 请同步简化 test_frontend_payload_dropped_bypass_keys"
    )

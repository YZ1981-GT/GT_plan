"""D4-4 营业收入调整分录汇总 导入导出往返守卫（纯函数级，不连库）。

背景（2026-09-28 修复）：`_SHEET_HEADERS["D4-4"]` 原为 **8 列**，漏了源模板
`营业收入调整分录汇总D4-4`（A1:J23，表头 R5）的 **F 列「……」**（前端字段 ``placeholder``）
与 **J 列「备注」**（前端字段 ``remark``）。后果不是报错而是**静默丢数据**：导出不带出这两列，
用户改后再导入时 `_parse_d4_4_row` 也不解析它们，前端 `safeParseRows` 对缺键兜底成空串
⇒ 用户填的内容无声消失。

本守卫钉死三层判据：
1. **列数/列序**对齐源模板 A~J 十列（并直接读权威册表头对账，防再次漏列）；
2. **export 与 headers 逐位等长**（通用不变量，错位/漏列都会被抓）；
3. **往返等值**：前端 10 个业务字段 export → parse 后逐字段一致，尤其 placeholder / remark。

断言的是行为（列数、列序、往返值），不是"函数存在"。
"""

from pathlib import Path

import pytest
from openpyxl import Workbook, load_workbook

from app.routers.wp_render_strategies import _d4_import_export as m

# 前端 `D4AdjustmentRow`（useD4Adjustment.ts）的 10 个业务字段（rowId 不参与导出）
_FRONTEND_BUSINESS_FIELDS = (
    "description",
    "category",
    "reportItem",
    "accountName",
    "noteItem",
    "placeholder",
    "debitAmount",
    "creditAmount",
    "indexRef",
    "remark",
)

# 源模板 R5 表头 A~J 十列（openpyxl 实测冻结）
_TEMPLATE_HEADER_ROW_5 = (
    "调整事项说明",
    "类别（报表调整/账项调整/其他）",
    "报表项目",
    "科目名称",
    "附注项目",
    "……",
    "借方调整金额",
    "贷方调整金额",
    "索引",
    "备注",
)

_REPO_ROOT = Path(__file__).resolve().parents[2]
# 权威册与 D4 各 provider 的 `TEMPLATE_RELATIVE_PATH` 保持一致（单一真源）。
# 注：`D4-1至D4-4 …（Leap-常规程序）.xlsx` 里的同名 sheet 几何逐字相同（已 openpyxl 实测），
# 但契约/instrumentation 锚定的是本册，故对账也锚定本册，避免将来两册分叉时判据失准。
_TEMPLATE_BOOK = _REPO_ROOT / "backend" / "wp_templates" / "D" / "D4 收入底稿.xlsx"
_TEMPLATE_SHEET = "营业收入调整分录汇总D4-4"


def _sample_row() -> dict:
    """一条填满的调整分录（10 业务字段全非空，便于抓"某列没带出"）。"""
    return {
        "rowId": "d4a-ms2p8tkl-juz5kck",
        "description": "跨期收入调整：12月发出商品未确认收入",
        "category": "账项调整",
        "reportItem": "营业收入",
        "accountName": "6001-主营业务收入",
        "noteItem": "营业收入及营业成本",
        "placeholder": "补充：已取得客户签收单",
        "debitAmount": 123456.78,
        "creditAmount": 0.0,
        "indexRef": "D4-17",
        "remark": "已与管理层沟通并同意调整",
    }


# ═══════════════════════════════════════════════════════════════════════════
# 一、列数 / 列序
# ═══════════════════════════════════════════════════════════════════════════


def test_d4_4_headers_have_ten_columns():
    """列头恰 10 列 —— 与源模板 A~J 一一对应（修复前为 8 列）。"""
    headers = m._SHEET_HEADERS["D4-4"]
    assert len(headers) == 10, f"D4-4 列头应为 10 列（源模板 A~J），实为 {len(headers)}：{headers}"


def test_d4_4_headers_cover_placeholder_and_remark():
    """F 列（……→补充说明）与 J 列（备注）必须在列头内 —— 这正是修复前漏掉的两列。"""
    headers = m._SHEET_HEADERS["D4-4"]
    assert "补充说明" in headers, "缺 F 列（源模板「……」/ 前端 placeholder）"
    assert "备注" in headers, "缺 J 列（源模板「备注」/ 前端 remark）"


def test_d4_4_header_order_matches_template_columns():
    """列序逐位对齐源模板 A~J（错位会让导入按名取值时串列）。"""
    headers = m._SHEET_HEADERS["D4-4"]
    expected = [
        "摘要",        # A 调整事项说明
        "分类",        # B 类别
        "报表项目",     # C
        "会计科目",     # D 科目名称
        "附注项目",     # E
        "补充说明",     # F ……
        "借方",        # G 借方调整金额
        "贷方",        # H 贷方调整金额
        "索引号",       # I 索引
        "备注",        # J
    ]
    assert list(headers) == expected


@pytest.mark.skipif(not _TEMPLATE_BOOK.exists(), reason="权威模板册不在工作树内")
def test_d4_4_header_count_matches_authoritative_workbook():
    """把列数钉死到**权威册实际表头**，而非仅对齐一份手抄常量。

    这是根治「语义投影头 ≠ 物理列头」的判据：源模板若将来加列，本断言先红。
    """
    wb = load_workbook(_TEMPLATE_BOOK, read_only=True)
    try:
        assert _TEMPLATE_SHEET in wb.sheetnames, f"权威册缺 sheet {_TEMPLATE_SHEET}"
        ws = wb[_TEMPLATE_SHEET]
        row5 = next(ws.iter_rows(min_row=5, max_row=5, values_only=True))
    finally:
        wb.close()
    template_cols = [str(v).strip() for v in row5 if v is not None and str(v).strip()]
    assert tuple(template_cols) == _TEMPLATE_HEADER_ROW_5, (
        f"源模板 R5 表头已变化，实测 {template_cols}"
    )
    assert len(m._SHEET_HEADERS["D4-4"]) == len(template_cols), (
        f"导出列数 {len(m._SHEET_HEADERS['D4-4'])} != 源模板物理列数 {len(template_cols)}"
    )


# ═══════════════════════════════════════════════════════════════════════════
# 二、export 与 headers 等长（通用不变量）
# ═══════════════════════════════════════════════════════════════════════════


def test_d4_4_export_row_length_equals_headers():
    """导出行长度必须等于列头长度 —— 漏列/多列都会在此被抓。"""
    headers = m._SHEET_HEADERS["D4-4"]
    values = m._export_d4_4_row(_sample_row())
    assert len(values) == len(headers), (
        f"导出 {len(values)} 个值 vs {len(headers)} 个列头：{list(zip(headers, values))}"
    )


def test_d4_4_export_tolerates_missing_and_none_values():
    """缺键 / None 值不得抛错，一律归一为空串或 0（对齐 _export_d4_1_row 范式）。"""
    values = m._export_d4_4_row({})
    assert len(values) == len(m._SHEET_HEADERS["D4-4"])
    assert values[5] == "", "placeholder 缺键应为空串"
    assert values[9] == "", "remark 缺键应为空串"
    assert values[6] == 0.0 and values[7] == 0.0, "金额缺键应为 0.0"

    none_values = m._export_d4_4_row({k: None for k in _FRONTEND_BUSINESS_FIELDS})
    assert none_values[5] == "" and none_values[9] == ""
    assert none_values[6] == 0.0 and none_values[7] == 0.0


# ═══════════════════════════════════════════════════════════════════════════
# 三、往返等值
# ═══════════════════════════════════════════════════════════════════════════


def _roundtrip(row: dict) -> dict:
    """export → 写 xlsx → 按实际列头读回 → parse，返回解析后的 dict。"""
    headers = list(m._SHEET_HEADERS["D4-4"])
    wb = Workbook()
    ws = wb.active
    ws.append(headers)
    ws.append(m._export_d4_4_row(row))
    all_rows = list(ws.iter_rows(values_only=True))
    actual_headers = [str(h).strip() if h is not None else "" for h in all_rows[0]]
    return m._parse_d4_4_row(all_rows[1], actual_headers, headers)


def test_d4_4_roundtrip_preserves_all_business_fields():
    """10 个业务字段 export → parse 后逐字段一致（rowId 由导入侧新生成，不比对）。"""
    src = _sample_row()
    got = _roundtrip(src)
    for field in _FRONTEND_BUSINESS_FIELDS:
        assert field in got, f"parse 结果缺字段 {field}（该列内容会被前端静默丢弃）"
        assert got[field] == src[field], f"字段 {field} 往返不一致：{src[field]!r} → {got[field]!r}"


def test_d4_4_roundtrip_preserves_placeholder_and_remark_specifically():
    """单独钉死修复前丢失的两列 —— 去掉任一侧映射，本用例即红。"""
    src = _sample_row()
    got = _roundtrip(src)
    assert got["placeholder"] == "补充：已取得客户签收单"
    assert got["remark"] == "已与管理层沟通并同意调整"


def test_d4_4_parse_result_covers_frontend_type():
    """parse 键集必须覆盖前端 D4AdjustmentRow 全部业务字段 + rowId。"""
    got = _roundtrip(_sample_row())
    missing = [f for f in _FRONTEND_BUSINESS_FIELDS if f not in got]
    assert not missing, f"parse 未覆盖前端字段：{missing}"
    assert got.get("rowId"), "导入侧应生成 rowId（行身份）"


def test_d4_4_roundtrip_empty_row_keeps_defaults():
    """空行往返：文本列为空串、金额为 0、category 回落默认「账项调整」。"""
    got = _roundtrip({k: "" for k in _FRONTEND_BUSINESS_FIELDS})
    assert got["category"] == "账项调整"
    assert got["placeholder"] == "" and got["remark"] == ""
    assert got["debitAmount"] == 0.0 and got["creditAmount"] == 0.0


def test_d4_4_still_registered_as_supported_sheet():
    """D4-4 仍在导入导出支持集内（本次修复不改变可用性）。"""
    assert "D4-4" in m._SUPPORTED_SHEETS


def test_d4_4_export_helper_is_actually_wired_into_export_route():
    """`_export_d4_4_row` 必须被导出路由真实调用 —— 防「抽了函数但调用点没接上」。

    本次把原内联的 `elif sheet == "D4-4": row_values = [...]` 抽成独立纯函数（对齐
    `_export_d4_1_row` / `_export_d4_34_*` 既有范式）。若只留函数、调用点仍是旧内联列表，
    上面所有往返用例依然全绿（它们直接调函数），却对真实导出毫无作用 —— 故此处断言接线。
    """
    src = Path(m.__file__).read_text(encoding="utf-8")
    assert "_export_d4_4_row(data_row)" in src, (
        "导出路由未调用 _export_d4_4_row —— 函数已抽出但调用点未接上"
    )
    # 同时确认旧内联实现未残留（否则两套并存，改一处不生效）
    assert 'data_row.get("description", "")' not in src, (
        "旧内联 D4-4 导出实现仍在，存在双源风险"
    )

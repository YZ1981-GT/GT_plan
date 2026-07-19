"""G6 其他债权投资(SPPI组) — 导入导出 + AI 端点集成测试.

Tests:
  - 4表导出模板端点 (G6-5/G6-6/G6-9/G6-10)
  - 导出数据(空工作簿)
  - 导入数据(有效xlsx / 无效格式 / 无效sheet)
  - AI 4 section端点 + 无效section
"""

from __future__ import annotations

import io
from unittest.mock import AsyncMock

import pytest
from httpx import ASGITransport, AsyncClient
from openpyxl import Workbook

from app.core.database import get_db
from app.deps import get_current_user
from app.main import app
from app.models.base import UserRole


class _FakeUser:
    id = "test-user-id"
    name = "Test User"
    email = "test@example.com"
    role = UserRole.admin


@pytest.fixture(autouse=True)
def override_deps():
    """Override DB and auth dependencies for all tests."""
    mock_db = AsyncMock()
    mock_db.commit = AsyncMock()

    # load_json_rows returns empty list (no rows in DB)
    mock_result = AsyncMock()
    mock_result.scalar_one_or_none = AsyncMock(return_value=None)
    mock_db.execute = AsyncMock(return_value=mock_result)

    async def _override_db():
        yield mock_db

    app.dependency_overrides[get_current_user] = lambda: _FakeUser()
    app.dependency_overrides[get_db] = _override_db
    yield mock_db
    app.dependency_overrides.clear()


def _make_xlsx_bytes(headers: list[str], rows: list[list]) -> bytes:
    """Create minimal xlsx in-memory with given headers and data rows."""
    wb = Workbook()
    ws = wb.active
    ws.append(headers)
    for row in rows:
        ws.append(row)
    buf = io.BytesIO()
    wb.save(buf)
    buf.seek(0)
    return buf.getvalue()


# ═══════════════════════════════════════════════════════════════════════════════
# Export Template Tests (4表)
# ═══════════════════════════════════════════════════════════════════════════════


@pytest.mark.asyncio
async def test_export_template_g6_5():
    transport = ASGITransport(app=app)
    async with AsyncClient(transport=transport, base_url="http://test") as client:
        resp = await client.post(
            "/api/workpapers/test-wp/g6-sppi/export-template?sheet=G6-5"
        )
    assert resp.status_code == 200
    assert "application/vnd.openxmlformats" in resp.headers["content-type"]


@pytest.mark.asyncio
async def test_export_template_g6_6():
    transport = ASGITransport(app=app)
    async with AsyncClient(transport=transport, base_url="http://test") as client:
        resp = await client.post(
            "/api/workpapers/test-wp/g6-sppi/export-template?sheet=G6-6"
        )
    assert resp.status_code == 200
    assert "application/vnd.openxmlformats" in resp.headers["content-type"]


@pytest.mark.asyncio
async def test_export_template_g6_9():
    transport = ASGITransport(app=app)
    async with AsyncClient(transport=transport, base_url="http://test") as client:
        resp = await client.post(
            "/api/workpapers/test-wp/g6-sppi/export-template?sheet=G6-9"
        )
    assert resp.status_code == 200
    assert "application/vnd.openxmlformats" in resp.headers["content-type"]


@pytest.mark.asyncio
async def test_export_template_g6_10():
    transport = ASGITransport(app=app)
    async with AsyncClient(transport=transport, base_url="http://test") as client:
        resp = await client.post(
            "/api/workpapers/test-wp/g6-sppi/export-template?sheet=G6-10"
        )
    assert resp.status_code == 200
    assert "application/vnd.openxmlformats" in resp.headers["content-type"]


# ═══════════════════════════════════════════════════════════════════════════════
# Export Data Test (empty workbook)
# ═══════════════════════════════════════════════════════════════════════════════


@pytest.mark.asyncio
async def test_export_data_empty():
    transport = ASGITransport(app=app)
    async with AsyncClient(transport=transport, base_url="http://test") as client:
        resp = await client.post(
            "/api/workpapers/test-wp/g6-sppi/export-data?sheet=G6-5"
        )
    assert resp.status_code == 200
    assert "application/vnd.openxmlformats" in resp.headers["content-type"]


# ═══════════════════════════════════════════════════════════════════════════════
# Import Data Tests
# ═══════════════════════════════════════════════════════════════════════════════


@pytest.mark.asyncio
async def test_import_data_g6_9():
    """Import valid xlsx with G6-9 data rows → 200, imported_count > 0."""
    xlsx_bytes = _make_xlsx_bytes(
        ["证券名称", "证券代码", "面值", "数量(盘点)", "数量(账面)"],
        [
            ["国开债2024A", "101001", 1000000, 100, 100],
            ["农发债2024B", "101002", 500000, 50, 50],
        ],
    )
    transport = ASGITransport(app=app)
    async with AsyncClient(transport=transport, base_url="http://test") as client:
        resp = await client.post(
            "/api/workpapers/test-wp/g6-sppi/import-data?sheet=G6-9",
            files={"file": ("test.xlsx", xlsx_bytes, "application/vnd.openxmlformats-officedocument.spreadsheetml.sheet")},
        )
    assert resp.status_code == 200
    data = resp.json()
    # ResponseWrapperMiddleware wraps as {code, message, data}
    payload = data.get("data", data)
    assert payload["imported_count"] > 0


@pytest.mark.asyncio
async def test_import_data_invalid_format():
    """Import non-xlsx (txt file) → 422."""
    transport = ASGITransport(app=app)
    async with AsyncClient(transport=transport, base_url="http://test") as client:
        resp = await client.post(
            "/api/workpapers/test-wp/g6-sppi/import-data?sheet=G6-9",
            files={"file": ("test.txt", b"invalid content", "text/plain")},
        )
    assert resp.status_code == 422


@pytest.mark.asyncio
async def test_import_data_invalid_sheet():
    """Import with unsupported sheet code → 400."""
    xlsx_bytes = _make_xlsx_bytes(["col1"], [["val"]])
    transport = ASGITransport(app=app)
    async with AsyncClient(transport=transport, base_url="http://test") as client:
        resp = await client.post(
            "/api/workpapers/test-wp/g6-sppi/import-data?sheet=G6-99",
            files={"file": ("test.xlsx", xlsx_bytes, "application/vnd.openxmlformats-officedocument.spreadsheetml.sheet")},
        )
    assert resp.status_code == 400


# ═══════════════════════════════════════════════════════════════════════════════
# AI Endpoint Tests (4 sections + invalid)
# ═══════════════════════════════════════════════════════════════════════════════


@pytest.mark.asyncio
async def test_ai_fair_value_conclusion(monkeypatch):
    async def _fake_chat(**_kwargs):
        return "公允价值测试审计结论初稿"

    monkeypatch.setattr(
        "app.routers.wp_render_strategies._g6_other_bond_investment_sppi_ai.chat_completion",
        _fake_chat,
    )
    transport = ASGITransport(app=app)
    async with AsyncClient(transport=transport, base_url="http://test") as client:
        resp = await client.post(
            "/api/workpapers/test-wp/g6-sppi/ai/fair-value-conclusion",
            json={"existingContent": "", "relatedContext": {}},
        )
    assert resp.status_code == 200
    payload = resp.json().get("data", resp.json())
    assert payload["content"] == "公允价值测试审计结论初稿"


@pytest.mark.asyncio
async def test_ai_interest_conclusion(monkeypatch):
    async def _fake_chat(**_kwargs):
        return "利息测算审计结论初稿"

    monkeypatch.setattr(
        "app.routers.wp_render_strategies._g6_other_bond_investment_sppi_ai.chat_completion",
        _fake_chat,
    )
    transport = ASGITransport(app=app)
    async with AsyncClient(transport=transport, base_url="http://test") as client:
        resp = await client.post(
            "/api/workpapers/test-wp/g6-sppi/ai/interest-conclusion",
            json={"existingContent": "", "relatedContext": {}},
        )
    assert resp.status_code == 200
    payload = resp.json().get("data", resp.json())
    assert payload["content"] == "利息测算审计结论初稿"


@pytest.mark.asyncio
async def test_ai_business_model_conclusion(monkeypatch):
    async def _fake_chat(**_kwargs):
        return "业务模式分析综合判断初稿"

    monkeypatch.setattr(
        "app.routers.wp_render_strategies._g6_other_bond_investment_sppi_ai.chat_completion",
        _fake_chat,
    )
    transport = ASGITransport(app=app)
    async with AsyncClient(transport=transport, base_url="http://test") as client:
        resp = await client.post(
            "/api/workpapers/test-wp/g6-sppi/ai/business-model-conclusion",
            json={"existingContent": "", "relatedContext": {}},
        )
    assert resp.status_code == 200
    payload = resp.json().get("data", resp.json())
    assert payload["content"] == "业务模式分析综合判断初稿"


@pytest.mark.asyncio
async def test_ai_sppi_conclusion(monkeypatch):
    async def _fake_chat(**_kwargs):
        return "SPPI测试综合结论初稿"

    monkeypatch.setattr(
        "app.routers.wp_render_strategies._g6_other_bond_investment_sppi_ai.chat_completion",
        _fake_chat,
    )
    transport = ASGITransport(app=app)
    async with AsyncClient(transport=transport, base_url="http://test") as client:
        resp = await client.post(
            "/api/workpapers/test-wp/g6-sppi/ai/sppi-conclusion",
            json={"existingContent": "", "relatedContext": {}},
        )
    assert resp.status_code == 200
    payload = resp.json().get("data", resp.json())
    assert payload["content"] == "SPPI测试综合结论初稿"


@pytest.mark.asyncio
async def test_ai_invalid_section():
    """Invalid section name → 400."""
    transport = ASGITransport(app=app)
    async with AsyncClient(transport=transport, base_url="http://test") as client:
        resp = await client.post(
            "/api/workpapers/test-wp/g6-sppi/ai/invalid-section",
            json={"existingContent": "", "relatedContext": {}},
        )
    assert resp.status_code == 400


# ═══════════════════════════════════════════════════════════════════════════════
# G6-6 nest helpers（导入保留元数据）
# ═══════════════════════════════════════════════════════════════════════════════


def test_nest_g6_6_preserves_conclusion_and_cross_validation():
    from app.routers.wp_render_strategies._g6_other_bond_investment_sppi_import_export import (
        _nest_g6_6_flat_rows,
    )

    nested = _nest_g6_6_flat_rows(
        [
            {
                "investProject": "债A",
                "faceValue": 100,
                "couponRate": 0.04,
                "effectiveRate": 0.05,
                "periodEnd": "2024-12-31",
                "openingAmortized": 98,
                "days": 365,
            }
        ],
        existing={
            "conclusion": "既有审计结论",
            "crossValidation": {
                "bookInterestIncome": 12.5,
                "interestAdjPeriodChange": 1.2,
                "auditedInterest": 1.2,
            },
        },
    )
    assert nested["conclusion"] == "既有审计结论"
    assert nested["crossValidation"]["bookInterestIncome"] == 12.5
    assert nested["crossValidation"]["interestAdjPeriodChange"] == 1.2
    assert nested["crossValidation"]["auditedInterest"] == 1.2
    assert len(nested["groups"]) == 1
    assert nested["groups"][0]["investProject"] == "债A"


def test_nest_g6_6_reuses_period_id_from_row_or_existing():
    from app.routers.wp_render_strategies._g6_other_bond_investment_sppi_import_export import (
        _flatten_g6_6_groups,
        _nest_g6_6_flat_rows,
    )

    existing = {
        "conclusion": "keep",
        "groups": [{
            "id": "g-stable",
            "investProject": "债A",
            "faceValue": 100,
            "couponRate": 0.04,
            "effectiveRate": 0.05,
            "periods": [{
                "id": "p-stable",
                "periodEnd": "2024-12-31",
                "openingAmortized": 98,
                "days": 365,
            }],
        }],
        "crossValidation": {},
    }
    # 有 periodId 时直接复用
    with_pid = _nest_g6_6_flat_rows(
        [{"investProject": "债A", "periodId": "p-from-flat", "periodEnd": "2024-06-30", "days": 180}],
        existing=existing,
    )
    assert with_pid["groups"][0]["periods"][0]["id"] == "p-from-flat"

    # Excel 无 periodId：按截止日匹配旧 id
    by_date = _nest_g6_6_flat_rows(
        [{"investProject": "债A", "periodEnd": "2024-12-31", "days": 365}],
        existing=existing,
    )
    assert by_date["groups"][0]["id"] == "g-stable"
    assert by_date["groups"][0]["periods"][0]["id"] == "p-stable"

    # flatten 写出 periodId
    flat = _flatten_g6_6_groups(existing["groups"])
    assert flat[0]["periodId"] == "p-stable"


def test_nest_g6_6_without_existing_defaults_meta():
    from app.routers.wp_render_strategies._g6_other_bond_investment_sppi_import_export import (
        _nest_g6_6_flat_rows,
    )

    nested = _nest_g6_6_flat_rows(
        [{"investProject": "债B", "faceValue": 1, "periodEnd": "2024-06-30", "days": 180}]
    )
    assert nested["conclusion"] == ""
    assert nested["crossValidation"]["bookInterestIncome"] == 0
    assert nested["crossValidation"]["interestAdjPeriodChange"] == 0


def test_nest_g6_6_merges_missing_fields_from_existing():
    """旧 11 列导入缺 stage/减值/初始确认时，从 existing 按截止日合并。"""
    from app.routers.wp_render_strategies._g6_other_bond_investment_sppi_import_export import (
        _flatten_g6_6_groups,
        _nest_g6_6_flat_rows,
    )

    existing = {
        "conclusion": "keep",
        "groups": [{
            "id": "g-stable",
            "crossSheetInvestmentId": "x-1",
            "investProject": "债A",
            "faceValue": 100,
            "couponRate": 0.04,
            "effectiveRate": 0.05,
            "purchasePrice": 98,
            "transactionCost": 1,
            "initialDate": "2023-01-01",
            "initialCarryingAmount": 99,
            "periods": [{
                "id": "p-stable",
                "periodStart": "2024-01-01",
                "periodEnd": "2024-12-31",
                "openingAmortized": 98,
                "openingImpairment": 5,
                "stage": "Stage3",
                "principalRecovered": 2,
                "days": 365,
                "daysManualOverride": True,
                "openingManualOverride": True,
                "indexRef": "IDX-1",
            }],
        }],
        "crossValidation": {},
    }
    # 模拟旧 Excel：仅核心 11 列字段
    nested = _nest_g6_6_flat_rows(
        [{
            "investProject": "债A",
            "faceValue": 100,
            "couponRate": 0.04,
            "effectiveRate": 0.05,
            "cutoffDate": "2024-12-31",
            "openingAmortized": 98,
            "effectiveInterest": 4,
            "cashInflow": 4,
            "endingAmortized": 98,
            "days": 365,
            "remark": "updated",
        }],
        existing=existing,
    )
    g = nested["groups"][0]
    p = g["periods"][0]
    assert g["id"] == "g-stable"
    assert g["purchasePrice"] == 98
    assert g["transactionCost"] == 1
    assert g["initialDate"] == "2023-01-01"
    assert g["crossSheetInvestmentId"] == "x-1"
    assert p["id"] == "p-stable"
    assert p["stage"] == "Stage3"
    assert p["openingImpairment"] == 5
    assert p["principalRecovered"] == 2
    assert p["periodStart"] == "2024-01-01"
    assert p["indexRef"] == "IDX-1"
    assert p["remark"] == "updated"
    assert p["openingManualOverride"] is True

    flat = _flatten_g6_6_groups(nested["groups"])
    assert "openingManualOverride" in flat[0]
    assert flat[0]["stage"] == "Stage3"


def test_g6_6_headers_include_extended_fields():
    from app.routers.wp_render_strategies._g6_other_bond_investment_sppi_import_export import (
        _G6_6_HEADERS,
        _G6_6_KEYS,
    )

    assert "减值阶段" in _G6_6_HEADERS
    assert "已收回本金" in _G6_6_HEADERS
    assert "起息日" in _G6_6_HEADERS
    assert "stage" in _G6_6_KEYS
    assert "principalRecovered" in _G6_6_KEYS
    assert "openingManualOverride" in _G6_6_KEYS
    assert "dayCountBasis" in _G6_6_KEYS
    assert "计息基准" in _G6_6_HEADERS
    assert len(_G6_6_HEADERS) == len(_G6_6_KEYS)


def test_g6_9_nest_flatten_preserves_meta_and_fields():
    """扁平 ↔ 嵌套往返：保留 meta，并展开差异/权属列。"""
    from app.routers.wp_render_strategies._g6_other_bond_investment_sppi_import_export import (
        _flatten_g6_9_items,
        _nest_g6_9_flat_rows,
    )

    existing = {
        "auditConclusion": "盘点结论",
        "meta": {"inventoryDate": "2024-12-31", "isBalanceSheetDate": True},
        "items": [],
    }
    flat_in = [
        {
            "id": "r1",
            "securitiesName": "国开债",
            "securitiesCode": "101001",
            "faceValue": 1000,
            "countQuantity": 110,
            "bookQuantity": 100,
            "varianceReason": "托管差异",
            "varianceConclusion": "拟调整",
            "ownershipEntity": "本公司",
            "restrictionType": "pledge",
            "restrictionNote": "质押给银行",
            "evidenceIndex": "E-1",
            "indexRef": "G6-9",
        }
    ]
    nested = _nest_g6_9_flat_rows(flat_in, existing=existing)
    assert nested["auditConclusion"] == "盘点结论"
    assert nested["meta"]["inventoryDate"] == "2024-12-31"
    assert nested["items"][0]["varianceReason"] == "托管差异"
    assert nested["items"][0]["restrictionType"] == "pledge"

    flat_out = _flatten_g6_9_items(nested["items"])
    assert flat_out[0]["securitiesCode"] == "101001"
    assert flat_out[0]["evidenceIndex"] == "E-1"
    assert flat_out[0]["ownershipEntity"] == "本公司"


@pytest.mark.asyncio
async def test_import_g6_9_dual_writes_primary_and_rows(monkeypatch, override_deps):
    """导入 G6-9 应双写嵌套主键 + 扁平 G6-9-rows。"""
    captured: list[tuple[str, object]] = []

    async def _fake_upsert(db, wp_id, item_id, payload, field="conclusion"):
        if field == "conclusion":
            captured.append((item_id, payload))

    async def _fake_load(*_a, **_k):
        return None

    monkeypatch.setattr(
        "app.routers.wp_render_strategies._g6_other_bond_investment_sppi_import_export.upsert_json_payload",
        _fake_upsert,
    )
    monkeypatch.setattr(
        "app.routers.wp_render_strategies._g6_other_bond_investment_sppi_import_export.load_json_payload",
        _fake_load,
    )

    xlsx_bytes = _make_xlsx_bytes(
        [
            "证券名称", "证券代码", "面值", "数量(盘点)", "数量(账面)",
            "差异原因", "差异结论", "权属主体", "受限类型", "受限说明", "证据索引", "索引",
        ],
        [
            ["国开债2024A", "101001", 1000000, 120, 100, "短少", "拟调整", "本公司", "none", "", "E1", "G6-9"],
        ],
    )
    transport = ASGITransport(app=app)
    async with AsyncClient(transport=transport, base_url="http://test") as client:
        resp = await client.post(
            "/api/workpapers/test-wp/g6-sppi/import-data?sheet=G6-9",
            files={
                "file": (
                    "test.xlsx",
                    xlsx_bytes,
                    "application/vnd.openxmlformats-officedocument.spreadsheetml.sheet",
                )
            },
        )
    assert resp.status_code == 200
    item_ids = [c[0] for c in captured]
    assert "G6-9-securities-inventory-data" in item_ids
    assert "G6-9-rows" in item_ids

    primary = next(p for iid, p in captured if iid == "G6-9-securities-inventory-data")
    rows = next(p for iid, p in captured if iid == "G6-9-rows")
    assert isinstance(primary, dict) and "items" in primary
    assert isinstance(rows, list) and len(rows) == 1
    assert primary["items"][0]["securitiesName"] == "国开债2024A"
    assert rows[0]["countQuantity"] == 120
    assert rows[0]["varianceConclusion"] == "拟调整"


def test_g6_10_nest_flatten_preserves_code_and_tx_types():
    from app.routers.wp_render_strategies._g6_other_bond_investment_sppi_import_export import (
        _flatten_g6_10_items,
        _nest_g6_10_flat_rows,
        _normalize_g6_10_tx_type,
        _validate_g6_10_nested,
    )

    assert _normalize_g6_10_tx_type("转入") == "transfer_in"
    assert _normalize_g6_10_tx_type("转出") == "transfer_out"
    assert _normalize_g6_10_tx_type("转让(转出)") == "transfer"

    nested = _nest_g6_10_flat_rows(
        [
            {
                "securitiesName": "国开债",
                "securitiesCode": "101001",
                "countDateQuantity": 100,
                "changeQuantity": 10,
                "bookQuantity": 90,
            },
            {
                "securitiesName": "国开债",
                "securitiesCode": "101001",
                "date": "2025-01-05",
                "transactionType": "转入",
                "quantity": 10,
                "voucherNo": "V1",
            },
        ]
    )
    assert nested["items"][0]["securitiesCode"] == "101001"
    assert nested["changeDetails"][0]["transactionType"] == "transfer_in"
    assert nested["changeDetails"][0]["securitiesCode"] == "101001"
    assert _validate_g6_10_nested(nested) == []

    flat = _flatten_g6_10_items(nested)
    assert any(r.get("securitiesCode") == "101001" and "countDateQuantity" in r for r in flat)
    assert any(r.get("transactionType") == "transfer_in" for r in flat)


def test_validate_g6_10_rejects_empty_illegal_and_orphan():
    from app.routers.wp_render_strategies._g6_other_bond_investment_sppi_import_export import (
        _validate_g6_10_nested,
    )

    assert any("为空" in e["reason"] for e in _validate_g6_10_nested({"items": [], "changeDetails": []}))

    errs = _validate_g6_10_nested({
        "items": [{"securitiesName": "债A"}],
        "changeDetails": [{"securitiesName": "债B", "transactionType": "回购"}],
    })
    reasons = [e["reason"] for e in errs]
    assert any("非法交易类型" in r for r in reasons)
    assert any("未出现在倒轧计算表" in r for r in reasons)


@pytest.mark.asyncio
async def test_import_g6_10_rejects_empty_overwrite(monkeypatch, override_deps):
    async def _fake_load(*_a, **_k):
        return {"items": [{"securitiesName": "已有"}], "changeDetails": []}

    monkeypatch.setattr(
        "app.routers.wp_render_strategies._g6_other_bond_investment_sppi_import_export.load_json_payload",
        _fake_load,
    )

    xlsx_bytes = _make_xlsx_bytes(["证券名称", "盘点日数量"], [])
    transport = ASGITransport(app=app)
    async with AsyncClient(transport=transport, base_url="http://test") as client:
        resp = await client.post(
            "/api/workpapers/test-wp/g6-sppi/import-data?sheet=G6-10",
            files={
                "file": (
                    "empty.xlsx",
                    xlsx_bytes,
                    "application/vnd.openxmlformats-officedocument.spreadsheetml.sheet",
                )
            },
        )
    assert resp.status_code == 422


def test_g6_10_multi_sheet_round_trip_preserves_code_and_transfer_in():
    """多 sheet 导出 → 解析 → nest → 再导出 → 再解析，保留代码与转入类型。"""
    from openpyxl import load_workbook

    from app.routers.wp_render_strategies._g6_other_bond_investment_sppi_import_export import (
        _build_g6_10_multi_sheet_workbook,
        _nest_g6_10_flat_rows,
        _parse_g6_10_import,
        _validate_g6_10_nested,
    )

    nested_in = {
        "activeTab": "rollForward",
        "selectedRowIndex": 0,
        "formulaVersion": 2,
        "changeConvention": "period_net_increase",
        "header": {"countDate": "2025-01-15", "balanceSheetDate": "2024-12-31"},
        "items": [
            {
                "id": "i1",
                "seq": 1,
                "securitiesName": "国开债",
                "securitiesCode": "101001",
                "countDateQuantity": 120,
                "changeQuantity": 20,
                "bookQuantity": 100,
                "varianceReason": "",
                "varianceConclusion": "",
                "indexRef": "G6-10",
                "remark": "",
            }
        ],
        "changeDetails": [
            {
                "id": "d1",
                "securitiesName": "国开债",
                "securitiesCode": "101001",
                "date": "2025-01-05",
                "transactionType": "transfer_in",
                "quantity": 20,
                "amount": 0,
                "voucherNo": "V1",
                "handler": "张三",
                "remark": "",
            }
        ],
    }
    assert _validate_g6_10_nested(nested_in) == []

    wb1 = _build_g6_10_multi_sheet_workbook(nested_in, template_only=False)
    buf1 = io.BytesIO()
    wb1.save(buf1)
    flat1, errors1 = _parse_g6_10_import(buf1.getvalue())
    assert errors1 == [] or all("缺少" not in e.get("reason", "") for e in errors1)
    nested2 = _nest_g6_10_flat_rows(flat1)
    assert nested2["items"][0]["securitiesCode"] == "101001"
    assert nested2["items"][0]["countDateQuantity"] == 120
    assert nested2["changeDetails"][0]["transactionType"] == "transfer_in"
    assert nested2["changeDetails"][0]["securitiesCode"] == "101001"
    assert nested2["changeDetails"][0]["quantity"] == 20

    wb2 = _build_g6_10_multi_sheet_workbook(nested2, template_only=False)
    buf2 = io.BytesIO()
    wb2.save(buf2)
    # 模板应含证券代码列
    wb_check = load_workbook(io.BytesIO(buf2.getvalue()), read_only=True)
    roll_name = next(n for n in wb_check.sheetnames if "倒轧" in n)
    headers = [c.value for c in next(wb_check[roll_name].iter_rows(min_row=2, max_row=2))]
    assert "证券代码" in headers
    wb_check.close()

    flat2, _ = _parse_g6_10_import(buf2.getvalue())
    nested3 = _nest_g6_10_flat_rows(flat2)
    assert nested3["items"][0]["securitiesCode"] == "101001"
    assert nested3["changeDetails"][0]["transactionType"] == "transfer_in"

"""test_h0_source_template_facts.py — H0 源模板事实守卫（xlsx 为唯一裁决者）.

spec: h0-confirmation-source-fidelity-and-linkage
  Requirements 12.1 / 12.2 / 12.3；Property 3 / 7 / 14 / 16 / 19 / 27 / 29

本文件是 H0 全部改造的**事实基线**：直读源模板 xlsx，并**直读前端 TS 源码**交叉锁死
（前端读不了 xlsx，故跨前后端比对统一在后端做 —— 这是平台既有范式）。

🔴 判 DV 必须 ``coord in dv.sqref`` 逐格测试：openpyxl 打印的 ``JF/JK/TG`` 等
远端列范围是 Excel 列重复残留，**不落在真实列上**。
"""
from __future__ import annotations

import json
import re
from pathlib import Path

import openpyxl
import pytest
from openpyxl.utils import get_column_letter

_REPO_ROOT = Path(__file__).resolve().parents[2]
H0_XLSX = _REPO_ROOT / "backend" / "wp_templates" / "H" / "H0 固定资产循环函证.xlsx"
_FE = _REPO_ROOT / "audit-platform" / "frontend" / "src" / "components" / "workpaper" / "confirmation"

COLUMN_SPEC_TS = _FE / "confirmationColumnSpec.ts"
MATRIX_TS = _FE / "h0SummaryMatrix.ts"
LOWER_TS = _FE / "h0SummaryLowerZone.ts"
FRAUD_TS = _FE / "fraudRisk" / "fraudRiskPresets.ts"
META_TS = _FE / "coordination" / "cycleConfirmationMeta.ts"

# ─── 源模板 9 张 sheet（全 visible） ─────────────────────────────────────────

EXPECTED_SHEETS = [
    "底稿目录",
    "函证程序表H0A",
    "函证结果汇总表H0-1",
    "核实被函证单位信息H0-2",
    "跟函函证过程控制H0-3",
    "差异核对表H0-4",
    "替代程序H0-5",
    "邮件传真回函可靠性验证H0-6",
    "函证程序舞弊风险评价表H0-7",
]

#: H0-1 上区 28 列 → 平台列 key（本表即证据表；改一侧必须同步另一侧）
SOURCE_COLUMN_TO_KEY: list[tuple[str, str, str]] = [
    ("A", "序号", "seq"),
    ("B", "询证函索引号", "confirm_index"),
    ("C", "选取样本目的", "sample_purpose"),
    ("D", "被询证单位名称", "entity_name"),
    ("E", "账户/交易", "account_type"),
    ("F", "金额或合同条款", "amount"),
    ("G", "函证方式", "send_channel"),
    ("H", "发函日期", "send_date"),
    ("I", "发函单号", "send_doc_no"),
    ("J", "收件地址", "entity_address"),
    ("K", "地址核查是否一致", "send_addr_match"),
    ("L", "是否收到回函", "is_replied"),
    ("M", "回函方式", "reply_method"),
    ("N", "是否相符", "match_status"),
    ("O", "回函日期", "reply_date"),
    ("P", "回函快递单号", "reply_courier_no"),
    ("Q", "回函发出地址", "reply_from_addr"),
    ("R", "发函地址与回函地址是否一致", "send_reply_addr_match"),
    ("S", "回函金额/条款", "reply_amount"),
    ("T", "差异（金额/条款）", "difference"),
    ("U", "可确认金额/条款", "confirmed_amount"),
    ("V", "差异核对索引（H0-4）", "diff_ref_index"),
    ("W", "其他说明/备注", "remark"),
    ("X", "是否采取替代程序", "use_alternative"),
    ("Y", "替代后可确认金额", "alt_confirmed"),
    ("Z", "替代后不可确认金额", "alt_unconfirmed"),
    ("AA", "替代程序索引号", "alt_ref_index"),
    ("AB", "审计结论", "row_conclusion"),
]

#: 源模板 H0-1 **没有**的三列（平台 BASE 有 → 必须剔除，否则空列噪声）
COLUMNS_ABSENT_IN_SOURCE = ["contact_person", "contact_phone", "currency"]

#: H0-1 上区五段合并表头（源模板 R5）
SOURCE_GROUP_HEADERS = [
    ("C5:F5", "发函询证纪要"),
    ("G5:K5", "1、发函信息"),
    ("L5:R5", "2、收到回函"),
    ("S5:W5", "3、回函金额确认"),
    ("X5:AA5", "4、未收到回函的替代程序"),
]

#: 七条 VLOOKUP（H0-1 ← H0-2），col_index_num 即源模板第三参
EXPECTED_VLOOKUP_COL_INDEX = {
    "D": 2,   # 被询证单位名称 ← H0-2!B
    "G": 3,   # 函证方式       ← H0-2!C
    "J": 4,   # 收件地址       ← H0-2!D
    "K": 10,  # 地址核查是否一致 ← H0-2!J
    "M": 16,  # 回函方式       ← H0-2!P
    "Q": 19,  # 回函发出地址   ← H0-2!S
    "R": 22,  # 发函地址与回函地址是否一致 ← H0-2!V
}

#: 矩阵 8 指标（源模板 C30:C37，去尾冒号）
EXPECTED_METRIC_LABELS = [
    "本期（期末）账面金额",
    "抽取样本的发函金额",
    "发函金额占账面金额的比例(%)",
    "回函确认金额",
    "回函可确认金额占发函金额的比例(%)",
    "回函可确认金额占账面金额的比例(%)",
    "替代测试确认金额",
    "回函和替代确认金额占账面金额的比例(%)",
]

#: 矩阵默认品种（源模板 E29/F29/G29；H29='……' 是可扩位不 seed）
EXPECTED_DEFAULT_CATEGORIES = ["固定资产", "工程物资", "租赁负债"]


# ─── fixtures / helpers ──────────────────────────────────────────────────────


@pytest.fixture(scope="module")
def wb():
    if not H0_XLSX.exists():
        pytest.fail(f"源模板不存在，无法裁决: {H0_XLSX}")
    return openpyxl.load_workbook(H0_XLSX)


@pytest.fixture(scope="module")
def wb_formula():
    return openpyxl.load_workbook(H0_XLSX, data_only=False)


def header_label(ws, col: str) -> str:
    """取该列表头：R6 有值取 R6，否则取 R5（A/B/AB 是 rowspan 合并）."""
    v6 = ws[f"{col}6"].value
    v5 = ws[f"{col}5"].value
    return str(v6 if v6 not in (None, "") else (v5 or "")).strip()


def dv_of(ws, coord: str) -> list[str]:
    return [dv.formula1 for dv in ws.data_validations.dataValidation if coord in dv.sqref]


def strip_ts_comments(src: str) -> str:
    return re.sub(r"/\*[\s\S]*?\*/", "", re.sub(r"(?m)(^|[^:])//[^\n]*", r"\1", src))


def ts_string_array(src: str, const_name: str) -> list[str]:
    """抽 `export const NAME = [ 'a', 'b', ] as const` 的字符串项（保序）."""
    m = re.search(rf"{const_name}\s*=\s*\[", src)
    assert m, f"未找到常量 {const_name}"
    start = src.index("[", m.start())
    depth = 0
    end = -1
    for i in range(start, len(src)):
        if src[i] == "[":
            depth += 1
        elif src[i] == "]":
            depth -= 1
            if depth == 0:
                end = i
                break
    assert end > start, f"{const_name} 括号未配对"
    return re.findall(r"'([^']*)'", src[start : end + 1])


def ts_block(src: str, anchor: str, open_ch: str = "{", close_ch: str = "}") -> str:
    """从 anchor **末尾**起做括号配对，取整块.

    🔴 必须从 anchor 末尾开始找 open_ch，不能从 anchor 开头找 ——
    TS 类型标注里常自带方括号（`ColumnDef[] = [` / `readonly H0SampleField[] = [`），
    从开头找会命中 `H0SampleField[` 的 `[` 并立刻与其 `]` 配对 → 得到空块 `[]`，
    断言退化为「解析出 0 条」（本守卫首版即因此把 12 条真实比对变成假失败）。
    故 anchor 约定**以 open_ch 结尾**。
    """
    assert anchor.endswith(open_ch), f"anchor 必须以 {open_ch!r} 结尾: {anchor!r}"
    m = src.index(anchor)
    start = m + len(anchor) - 1
    assert src[start] == open_ch
    depth = 0
    for i in range(start, len(src)):
        if src[i] == open_ch:
            depth += 1
        elif src[i] == close_ch:
            depth -= 1
            if depth == 0:
                return src[start : i + 1]
    raise AssertionError(f"{anchor} 括号未配对")


def test_ts_block_handles_type_annotation_brackets():
    """反向自检：anchor 里带类型方括号时仍能取到真数组块（首版 bug 的回归防护）."""
    sample = "export const X: readonly Foo[] = [\n  { a: 1 },\n]\n"
    assert ts_block(sample, "X: readonly Foo[] = [", "[", "]").strip() == "[\n  { a: 1 },\n]".strip()


# ─── Property 29: 9 张 sheet 全 visible ─────────────────────────────────────


def test_nine_sheets_all_visible(wb):
    """H0 九张 sheet 名逐字 + 全部 visible（与 E0 的 10 张隐藏表情形不同）."""
    assert wb.sheetnames == EXPECTED_SHEETS
    for name in EXPECTED_SHEETS:
        assert wb[name].sheet_state == "visible", f"{name} 不是 visible"


def test_no_adjudication_sheet(wb):
    """`审定表H0-1` **不存在** —— 公式预设的 sheet 名贴错标签（R11.1 的依据）."""
    assert "审定表H0-1" not in wb.sheetnames
    assert "函证结果汇总表H0-1" in wb.sheetnames


def test_index_sheet_lists_eight_workpapers(wb):
    """底稿目录 8 项索引（H0A / H0-1~H0-7）."""
    ws = wb["底稿目录"]
    pairs = [(ws.cell(r, 5).value, ws.cell(r, 6).value) for r in range(4, 12)]
    assert pairs == [
        ("函证程序表", "H0A"),
        ("函证结果汇总表", "H0-1"),
        ("核实被函证单位信息", "H0-2"),
        ("跟函函证过程控制", "H0-3"),
        ("差异核对表", "H0-4"),
        ("替代程序", "H0-5"),
        ("邮件传真回函可靠性验证", "H0-6"),
        ("函证程序舞弊风险评价表", "H0-7"),
    ]


def test_cycle_meta_matches_index_sheet():
    """`cycleConfirmationMeta.H0` 的 7 个索引号与底稿目录一致."""
    src = strip_ts_comments(META_TS.read_text(encoding="utf-8"))
    block = ts_block(src, "H0: {")
    for field, code in [
        ("summaryCode", "H0-1"),
        ("entityVerifyCode", "H0-2"),
        ("followupCode", "H0-3"),
        ("diffCode", "H0-4"),
        ("altPrimaryCode", "H0-5"),
        ("reliabilityCode", "H0-6"),
        ("fraudCode", "H0-7"),
    ]:
        assert re.search(rf"{field}:\s*'{re.escape(code)}'", block), f"H0 meta {field} != {code}"


# ─── Property 16: H0-1 28 列表头逐字 + 与前端 label 交叉锁死 ─────────────────


def test_h0_1_has_exactly_28_columns(wb):
    ws = wb["函证结果汇总表H0-1"]
    labels = [header_label(ws, get_column_letter(c)) for c in range(1, 29)]
    assert all(labels), f"A..AB 应全部有表头，实得: {labels}"
    assert labels == [lbl for _, lbl, _ in SOURCE_COLUMN_TO_KEY]
    # 第 29 列（AC）起不应再有表头
    assert not header_label(ws, "AC")


def test_h0_1_group_headers(wb):
    """五段合并表头逐字（源模板 R5 的合并区）."""
    ws = wb["函证结果汇总表H0-1"]
    merged = {str(r) for r in ws.merged_cells.ranges}
    for rng, label in SOURCE_GROUP_HEADERS:
        assert rng in merged, f"合并区 {rng} 不存在"
        assert str(ws[rng.split(":")[0]].value).strip() == label
    # A/B/AB 是 rowspan=3 的独立列（不属任何段）
    for rng in ("A5:A7", "B5:B7", "AB5:AB7"):
        assert rng in merged, f"{rng} 应为 rowspan 合并"


@pytest.fixture(scope="module")
def fe_column_labels() -> dict[str, str]:
    """前端 H0 的**有效** label = CYCLE_COLUMN_LABEL_OVERRIDES.H0 优先，否则 BASE/VARIANT."""
    src = strip_ts_comments(COLUMN_SPEC_TS.read_text(encoding="utf-8"))
    base = ts_block(src, "BASE_CONFIRMATION_COLUMNS: ColumnDef[] = [", "[", "]")
    variant = ts_block(src, "VARIANT_COLUMN_DEFS: Record<string, ColumnDef> = {")
    overrides = ts_block(src, "H0: Object.freeze({", "{", "}")

    labels: dict[str, str] = {}
    for blk in (base, variant):
        for key, label in re.findall(r"key:\s*'([^']+)',\s*label:\s*'([^']*)'", blk):
            labels.setdefault(key, label)
    for key, label in re.findall(r"(\w+):\s*'([^']*)'", overrides):
        labels[key] = label
    assert len(labels) > 25, f"前端 label 解析异常，只得 {len(labels)} 条"
    return labels


@pytest.mark.parametrize("col,src_label,key", SOURCE_COLUMN_TO_KEY)
def test_frontend_label_matches_source(col: str, src_label: str, key: str, fe_column_labels):
    """H0 每列的前端有效 label 与源模板表头逐字一致（R6.4）."""
    assert key in fe_column_labels, f"前端未定义列 key `{key}`（源模板 {col} 列 {src_label}）"
    assert fe_column_labels[key] == src_label, (
        f"源模板 {col} 列「{src_label}」 vs 前端 `{key}` label「{fe_column_labels[key]}」"
    )


def test_absent_columns_are_excluded_in_frontend():
    """源模板 H0-1 没有的三列必须在 CYCLE_EXCLUDED_COLUMNS.H0 里（R6.2）."""
    src = strip_ts_comments(COLUMN_SPEC_TS.read_text(encoding="utf-8"))
    m = re.search(r"CYCLE_EXCLUDED_COLUMNS[\s\S]*?H0:\s*\[([^\]]*)\]", src)
    assert m, "未找到 CYCLE_EXCLUDED_COLUMNS.H0"
    excluded = re.findall(r"'([^']+)'", m.group(1))
    assert sorted(excluded) == sorted(COLUMNS_ABSENT_IN_SOURCE)


def test_source_has_no_contact_or_currency_columns(wb):
    """反向自检：源模板 H0-1 表头确实不含联系人/联系电话/币种."""
    ws = wb["函证结果汇总表H0-1"]
    labels = {header_label(ws, get_column_letter(c)) for c in range(1, 29)}
    for absent in ("联系人", "联系电话", "币种"):
        assert absent not in labels, f"源模板 H0-1 竟含「{absent}」列，需重新评估剔除决策"
    # 联系人/联系电话在 H0-2（不是没有，是不在这张表）
    w2 = wb["核实被函证单位信息H0-2"]
    r6 = {str(w2.cell(6, c).value or "").strip() for c in range(1, 41)}
    assert "联系人" in r6 and "联系电话" in r6


# ─── Property 14: 七条 VLOOKUP 的 col_index_num ─────────────────────────────


def test_vlookup_col_index_num(wb_formula):
    """H0-1 ← H0-2 的七条 VLOOKUP 第三参逐条（R5.1/R5.6 的事实基线）."""
    ws = wb_formula["函证结果汇总表H0-1"]
    got: dict[str, int] = {}
    for col in EXPECTED_VLOOKUP_COL_INDEX:
        f = ws[f"{col}8"].value
        assert isinstance(f, str) and f.startswith("=VLOOKUP("), f"{col}8 不是 VLOOKUP: {f!r}"
        assert "核实被函证单位信息H0-2" in f, f"{col}8 未引用 H0-2: {f}"
        m = re.search(r",\s*(\d+)\s*,\s*0\s*\)", f)
        assert m, f"{col}8 未解析出 col_index_num: {f}"
        got[col] = int(m.group(1))
    assert got == EXPECTED_VLOOKUP_COL_INDEX


def test_vlookup_lookup_value_is_confirm_index(wb_formula):
    """lookup_value 是 B 列（询证函索引号）→ 带入匹配键必须用索引号（R5.2）."""
    ws = wb_formula["函证结果汇总表H0-1"]
    for col in EXPECTED_VLOOKUP_COL_INDEX:
        f = ws[f"{col}8"].value
        assert f.startswith("=VLOOKUP(B8,"), f"{col}8 的 lookup_value 不是 B8: {f}"


def test_vlookup_target_columns_semantics(wb_formula, wb):
    """col_index_num 指向的 H0-2 列语义与 H0-1 目标列一致（防映射错位）."""
    w2 = wb["核实被函证单位信息H0-2"]
    expected_pairs = {
        "D": "被询证单位全称",
        "G": "函证方式",
        "J": "收件地址",
        "K": "发函地址与企查查地址是否一致",
        "M": "回函方式",
        "Q": "回函发出地址（含物流信息中的地址）",
        "R": "发函地址与回函地址是否一致",
    }
    for col, idx in EXPECTED_VLOOKUP_COL_INDEX.items():
        h02_label = str(w2.cell(6, idx).value or "").strip()
        assert h02_label == expected_pairs[col], (
            f"H0-1 {col} 列 ← H0-2 第 {idx} 列，期望「{expected_pairs[col]}」实得「{h02_label}」"
        )


def test_frontend_pull_mapping_matches_vlookup(wb, wb_formula):
    """前端 `H0_PULL_FROM_ENTITY_VERIFY` 与源模板七条 VLOOKUP 逐条交叉锁死（Property 14）.

    比对三项：`sourceColumnIndex` == col_index_num ·
    `sourceColumn` 字母与 index 对应 · `label` 与 H0-1 表头一致。
    """
    pull_ts = _FE / "h0SummaryFromEntityVerify.ts"
    assert pull_ts.exists(), "缺少 h0SummaryFromEntityVerify.ts（Task 15）"
    src = strip_ts_comments(pull_ts.read_text(encoding="utf-8"))
    block = ts_block(src, "H0_PULL_FROM_ENTITY_VERIFY: readonly H0PullFieldSpec[] = [", "[", "]")
    entries = re.findall(r"\{([^{}]*)\}", block)
    assert len(entries) == 7, f"带入映射应为 7 条，实得 {len(entries)}"

    ws1 = wb["函证结果汇总表H0-1"]
    ws1f = wb_formula["函证结果汇总表H0-1"]
    ws2 = wb["核实被函证单位信息H0-2"]

    seen: dict[str, int] = {}
    for body in entries:
        target_col = re.search(r"targetColumn:\s*'([^']+)'", body).group(1)
        source_col = re.search(r"sourceColumn:\s*'([^']+)'", body).group(1)
        idx = int(re.search(r"sourceColumnIndex:\s*(\d+)", body).group(1))
        label = re.search(r"label:\s*'([^']*)'", body).group(1)
        seen[target_col] = idx

        # ① col_index_num 与源模板 VLOOKUP 第三参一致
        formula = ws1f[f"{target_col}8"].value
        assert isinstance(formula, str) and formula.startswith("=VLOOKUP(B8,"), (
            f"H0-1 {target_col}8 不是从 B 列 VLOOKUP: {formula!r}"
        )
        m = re.search(r",\s*(\d+)\s*,\s*0\s*\)", formula)
        assert m and int(m.group(1)) == idx, (
            f"{target_col} 列 col_index_num：源模板={m.group(1) if m else '?'} 前端={idx}"
        )

        # ② sourceColumn 字母与 index 指向同一列
        assert get_column_letter(idx) == source_col, (
            f"{target_col} 列：index {idx} 对应 {get_column_letter(idx)} 列，前端写 {source_col}"
        )
        assert str(ws2.cell(6, idx).value or "").strip(), f"H0-2 第 {idx} 列无表头"

        # ③ label 与 H0-1 表头一致
        assert header_label(ws1, target_col) == label, (
            f"{target_col} 列 label：源模板「{header_label(ws1, target_col)}」前端「{label}」"
        )

    assert seen == EXPECTED_VLOOKUP_COL_INDEX


def test_frontend_pull_targets_send_channel_not_confirmation_method():
    """「函证方式」列的带入目标必须是 `send_channel`（渠道），不是 `confirmation_method`.

    🔴 源模板 H0-1!G ← H0-2!C（DV = 邮寄/跟函/电子函证/其他）是**渠道**；
    `confirmation_method` 承载积极式/消极式并驱动可确认金额派生，
    往它写渠道值会让「消极式未回函→视同相符」分支永久失效。
    """
    src = strip_ts_comments((_FE / "h0SummaryFromEntityVerify.ts").read_text(encoding="utf-8"))
    block = ts_block(src, "H0_PULL_FROM_ENTITY_VERIFY: readonly H0PullFieldSpec[] = [", "[", "]")
    targets = re.findall(r"targetKey:\s*'([^']+)'", block)
    assert "send_channel" in targets
    assert "confirmation_method" not in targets, "带入不得写 confirmation_method（渠道 ≠ 积极式/消极式）"


def test_difference_formula(wb_formula):
    """T 列差异派生公式：`IF(L="是",S-F,"未回函")`."""
    ws = wb_formula["函证结果汇总表H0-1"]
    assert ws["T8"].value == '=IF(L8="是",S8-F8,"未回函")'


# ─── Property 3: 矩阵 8 指标 + 默认品种 ─────────────────────────────────────


def test_matrix_metric_labels(wb):
    """8 指标标签逐字（源模板 C30:C37，去尾冒号）."""
    ws = wb["函证结果汇总表H0-1"]
    got = [str(ws.cell(r, 3).value or "").strip().rstrip("：") for r in range(30, 38)]
    assert got == EXPECTED_METRIC_LABELS


def test_matrix_metric_labels_match_frontend():
    src = strip_ts_comments(MATRIX_TS.read_text(encoding="utf-8"))
    assert ts_string_array(src, "H0_MATRIX_METRIC_LABELS") == EXPECTED_METRIC_LABELS


def test_matrix_default_categories(wb):
    """默认品种 = E29/F29/G29；H29 是 `……` 可扩位（不 seed）."""
    ws = wb["函证结果汇总表H0-1"]
    assert [str(ws.cell(29, c).value or "").strip() for c in (5, 6, 7)] == EXPECTED_DEFAULT_CATEGORIES
    assert str(ws.cell(29, 8).value or "").strip() == "……"
    assert str(ws.cell(29, 3).value or "").strip() == "项目"


def test_matrix_default_categories_match_frontend():
    src = strip_ts_comments(MATRIX_TS.read_text(encoding="utf-8"))
    assert ts_string_array(src, "H0_MATRIX_DEFAULT_CATEGORIES") == EXPECTED_DEFAULT_CATEGORIES


def test_source_matrix_omits_right_of_use(wb):
    """反向自检：源模板矩阵**未列**「使用权资产」（源模板自身遗漏，登记不修）.

    H0A/H0-1/H0-4/H0-5 的标题都写「固定资产/工程物资/使用权资产/租赁负债/……」，
    但矩阵只有 3 列 → 故不 seed 使用权资产，由用户从枚举增列。
    """
    ws = wb["函证结果汇总表H0-1"]
    cats = [str(ws.cell(29, c).value or "").strip() for c in range(5, 12)]
    assert "使用权资产" not in cats
    assert "使用权资产" in str(ws["A2"].value)


def test_matrix_sumif_formulas(wb_formula):
    """三个聚合行按 SUMIF(E,品种,F/U/Y)；五个比例行按源模板分子分母."""
    ws = wb_formula["函证结果汇总表H0-1"]
    assert ws["E31"].value == "=SUMIF(E8:E27,E29,F8:F27)"
    assert ws["E33"].value == "=SUMIF(E8:E27,E29,U8:U27)"
    assert ws["E36"].value == "=SUMIF(E8:E27,E29,Y8:Y27)"
    assert ws["E32"].value == "=IF(ISERROR(E31/E30),0,E31/E30)"
    assert ws["E34"].value == "=IF(ISERROR(E33/E31),0,E33/E31)"
    assert ws["E35"].value == "=IF(ISERROR(E33/E30),0,E33/E30)"
    assert ws["E37"].value == "=(E36+E33)/E30"
    # R30 是手填（无公式）
    assert ws["E30"].value is None or not str(ws["E30"].value).startswith("=")


# ─── Property 7: 下区四块文字逐字 + 前端常量交叉锁死 ────────────────────────


def test_lower_zone_block_anchors(wb):
    """四块锚点与原文（🔴 S28 原文是「二、审计说明」= 源模板编号笔误）."""
    ws = wb["函证结果汇总表H0-1"]
    assert str(ws["C28"].value).strip() == "一、函证情况"
    assert str(ws["J28"].value).strip() == "二、样本选择"
    assert str(ws["S28"].value).strip() == "二、审计说明"  # 笔误原文
    assert str(ws["C39"].value).strip() == "四、审计结论"


def test_frontend_registers_s28_typo():
    """前端 `audit_note` 显示「三、审计说明」但登记 sourceText 为原文."""
    src = strip_ts_comments(LOWER_TS.read_text(encoding="utf-8"))
    block = ts_block(src, "H0_LOWER_ZONE_BLOCKS = {")
    assert re.search(r"audit_note:\s*\{[^}]*title:\s*'三、审计说明'", block)
    assert re.search(r"audit_note:\s*\{[^}]*sourceText:\s*'二、审计说明'", block)


def test_sample_selection_fields(wb):
    """二、样本选择 6 字段标签与锚点逐字（J29/J30/J31/J32/J34/J35）."""
    ws = wb["函证结果汇总表H0-1"]
    expected = [
        ("J29", "测试总体"),
        ("J30", "特定样本"),
        ("J31", "抽样总体"),
        ("J32", "确定的抽样样本量"),
        ("J34", "抽样方法"),
        ("J35", "抽样过程"),
    ]
    for anchor, label in expected:
        assert str(ws[anchor].value or "").strip().rstrip("：") == label, anchor
    # J33 / J36 / J37 无字段标签（源模板留空）
    for blank in ("J33", "J36", "J37"):
        assert not str(ws[blank].value or "").strip(), blank


def test_sample_fields_match_frontend(wb):
    """前端 6 字段的 label / anchor / placeholder 与源模板逐字（含跨格合并）."""
    ws = wb["函证结果汇总表H0-1"]
    src = strip_ts_comments(LOWER_TS.read_text(encoding="utf-8"))
    block = ts_block(src, "H0_SAMPLE_SELECTION_FIELDS: readonly H0SampleField[] = [", "[", "]")
    items = re.findall(
        r"key:\s*'([^']+)',\s*label:\s*'([^']*)',\s*anchor:\s*'([^']+)',[\s\S]*?"
        r"placeholder:\s*'([^']*)',\s*placeholderAnchor:\s*'([^']+)'",
        block,
    )
    assert len(items) == 6, f"前端样本选择字段解析出 {len(items)} 条"
    for _key, label, anchor, placeholder, ph_anchor in items:
        assert str(ws[anchor].value or "").strip().rstrip("：") == label, f"{anchor} label 不符"
        merged = "".join(str(ws[a].value or "") for a in ph_anchor.split("+"))
        assert placeholder == merged, (
            f"{ph_anchor} placeholder 与源模板不符（跨格合并）:\n  前端={placeholder!r}\n  源模板={merged!r}"
        )


def test_audit_note_sections(wb):
    """三、审计说明 5 小节标题与锚点逐字（S29/W29/S33/S34/S37）."""
    ws = wb["函证结果汇总表H0-1"]
    expected = [
        ("S29", "1、对询证函保持的控制的说明"),
        ("W29", "2、对误差的分析"),
        ("S33", "3、对以传真或电子邮件形式收到的回函的可靠性的考虑（H0-6）"),
        ("S34", "4、针对不符事项的程序"),
        ("S37", "5、针对未回函的替代程序"),
    ]
    for anchor, title in expected:
        assert str(ws[anchor].value or "").strip() == title, anchor


def test_audit_note_sections_match_frontend(wb):
    """前端 5 小节 title/anchor/hint 与源模板逐字（hint 含跨格合并）."""
    ws = wb["函证结果汇总表H0-1"]
    src = strip_ts_comments(LOWER_TS.read_text(encoding="utf-8"))
    block = ts_block(src, "H0_AUDIT_NOTE_SECTIONS: readonly H0AuditNoteSection[] = [", "[", "]")
    entries = re.findall(r"\{([^{}]*)\}", block)
    assert len(entries) == 5, f"前端审计说明小节解析出 {len(entries)} 条"
    for body in entries:
        title = re.search(r"title:\s*'([^']*)'", body).group(1)
        anchor = re.search(r"anchor:\s*'([^']+)'", body).group(1)
        assert str(ws[anchor].value or "").strip() == title, f"{anchor} title 不符"
        hm = re.search(r"hint:\s*'([^']*)'", body)
        ham = re.search(r"hintAnchor:\s*'([^']+)'", body)
        if hm:
            assert ham, f"{anchor} 有 hint 但缺 hintAnchor"
            merged = "".join(str(ws[a].value or "") for a in ham.group(1).split("+"))
            assert hm.group(1) == merged, (
                f"{ham.group(1)} hint 与源模板不符:\n  前端={hm.group(1)!r}\n  源模板={merged!r}"
            )


def test_error_condition_is_split_across_two_cells(wb):
    """反向自检：误差构成条件确实被源模板拆成 W30+W31（否则合并逻辑无意义）."""
    ws = wb["函证结果汇总表H0-1"]
    w30 = str(ws["W30"].value or "")
    w31 = str(ws["W31"].value or "")
    assert w30 and w31, "W30/W31 应都有内容"
    assert w30.endswith("人民币"), f"W30 应以「人民币」结尾（半句话）: {w30!r}"
    assert w31.startswith("（）万元"), f"W31 应以「（）万元」开头: {w31!r}"


def test_reference_conclusions(wb):
    """参考结论 A/B/C 逐字（A66:B68）+ 与前端一致."""
    ws = wb["函证结果汇总表H0-1"]
    src = strip_ts_comments(LOWER_TS.read_text(encoding="utf-8"))
    block = ts_block(src, "H0_REFERENCE_CONCLUSIONS: readonly { code: string; text: string; anchor: string }[] = [", "[", "]")
    items = re.findall(r"code:\s*'([A-C])',\s*text:\s*'([^']*)',\s*anchor:\s*'([^']+)'", block)
    assert len(items) == 3
    for code, text, anchor in items:
        a_ref, b_ref = anchor.split("+")
        assert str(ws[a_ref].value or "").strip().rstrip("、") == code, a_ref
        assert str(ws[b_ref].value or "").strip() == text, b_ref


def test_guidance_texts_match_frontend(wb):
    """编制说明（只读方法论上下文）每段与源模板逐格拼接一致."""
    ws = wb["函证结果汇总表H0-1"]
    src = strip_ts_comments(LOWER_TS.read_text(encoding="utf-8"))
    block = ts_block(src, "H0_LOWER_ZONE_TEXTS: readonly H0LowerZoneText[] = [", "[", "]")
    entries = re.findall(r"\{([^{}]*)\}", block)
    assert len(entries) >= 5, f"编制说明解析出 {len(entries)} 条"
    checked = 0
    for body in entries:
        am = re.search(r"anchor:\s*'([^']+)'", body)
        if not am:
            continue
        anchors = am.group(1).split("+")
        # text 可能是多行字符串拼接（'a\n' + 'b'），逐段取出后拼接
        parts = re.findall(r"'((?:[^'\\]|\\.)*)'", body.split("anchor:")[0])
        text = "".join(parts).replace("\\n", "\n").replace("\\u3000", "\u3000")
        src_text = "\n".join(str(ws[a].value or "") for a in anchors)
        # 去掉 key/block 值造成的噪声段（只比对最长的那段）
        assert text.strip(), f"{am.group(1)} 前端文字为空"
        assert src_text.strip(), f"{am.group(1)} 源模板文字为空"
        # 逐锚点校验：前端文字必须含每个源模板格的内容
        for a in anchors:
            cell = str(ws[a].value or "").strip()
            if cell:
                assert cell in text, f"锚点 {a} 的源模板内容未出现在前端文字中: {cell[:40]!r}"
        checked += 1
    assert checked >= 5, f"只校验了 {checked} 段编制说明"


# ─── Property 19: 五处真实 DV（逐格判定） ───────────────────────────────────


def test_h0_1_real_dv_only_four_columns(wb):
    """H0-1 真实 DV 只有 C（选样目的）/ L / N / X（是否类）四列."""
    ws = wb["函证结果汇总表H0-1"]
    with_dv = [
        get_column_letter(c)
        for c in range(1, 29)
        if dv_of(ws, f"{get_column_letter(c)}8")
    ]
    assert with_dv == ["C", "L", "N", "X"], f"H0-1 带 DV 的列: {with_dv}"


def test_mirror_column_dv_not_on_real_columns(wb):
    """反向自检：`跟函,邮寄,电邮,其他` 这条 DV 只落在镜像列上，不覆盖任何真实列."""
    ws = wb["函证结果汇总表H0-1"]
    target = '"跟函,邮寄,电邮,其他"'
    all_formulas = {dv.formula1 for dv in ws.data_validations.dataValidation}
    assert target in all_formulas, "源模板应含这条镜像列 DV；若消失需重写本自检"
    for c in range(1, 29):
        coord = f"{get_column_letter(c)}8"
        assert target not in dv_of(ws, coord), f"{coord} 竟命中镜像 DV"


def test_h0_2_dv_values(wb):
    ws = wb["核实被函证单位信息H0-2"]
    assert dv_of(ws, "C7") == ['"邮寄,跟函,电子函证,其他"']
    assert dv_of(ws, "P7") == ['"纸质原件,电子函证,其他介质"']
    assert dv_of(ws, "L7") == [
        '"发票/合同地址核实,电话核实,官网/公告查询,地图查询,邮件确认,其他方式"'
    ]
    assert dv_of(ws, "AB7") == ['"送抵,退回"']
    assert dv_of(ws, "AL7") == ['"送抵,退回"']


def test_h0_1_sample_purpose_dv(wb):
    ws = wb["函证结果汇总表H0-1"]
    assert dv_of(ws, "C8") == ['"A. 大额,B.异常,C.余额为0,D.账龄长,E.随机"']


def test_h0_6_dv(wb):
    ws = wb["邮件传真回函可靠性验证H0-6"]
    assert dv_of(ws, "D7") == ['"传真,电子邮件"']
    for coord in ("E7", "F7", "L7"):
        assert dv_of(ws, coord) == ['"是,否"'], coord


def test_h0_3_checkpoints_dv(wb):
    ws = wb["跟函函证过程控制H0-3"]
    for coord in ("B23", "C23"):
        assert dv_of(ws, coord) == ['"是,否"'], coord


def test_h0_4_has_no_dv(wb):
    """H0-4 源模板无 DV（差异表全手工录入）."""
    ws = wb["差异核对表H0-4"]
    for c in range(1, 10):
        assert not dv_of(ws, f"{get_column_letter(c)}6")


# ─── H0A 程序表（11 条，源模板漏编号 4） ────────────────────────────────────


def test_h0a_has_eleven_programs_with_source_numbering_gap(wb):
    """H0A 11 条程序，编号 1,2,3,**5**,6…12 —— 漏 4 属**源模板事实**，不得补."""
    ws = wb["函证程序表H0A"]
    nums = [str(ws.cell(r, 1).value or "").strip() for r in range(7, 18)]
    assert nums == ["1", "2", "3", "5", "6", "7", "8", "9", "10", "11", "12"]
    assert not str(ws.cell(18, 1).value or "").strip(), "第 12 行后不应还有程序"
    assert "4" not in nums, "源模板确实漏编号 4"


def test_h0a_program_categories(wb):
    """程序分类只有「常规★」与「备选」；第 7/8 条为备选."""
    ws = wb["函证程序表H0A"]
    cats = [str(ws.cell(r, 4).value or "").strip() for r in range(7, 18)]
    assert set(cats) == {"常规★", "备选"}
    assert cats == ["常规★"] * 5 + ["备选", "备选"] + ["常规★"] * 4
    assert str(ws["D5"].value).strip() == "程序分类"
    assert str(ws["E5"].value).strip() == "底稿索引号"


def test_h0a_ref_index_targets_exist(wb):
    """底稿索引号列引用的编码都在底稿目录里（第 8/12 条源模板未填，属事实）."""
    ws = wb["函证程序表H0A"]
    valid = {"H0A", "H0-1", "H0-2", "H0-3", "H0-4", "H0-5", "H0-6", "H0-7"}
    empty_rows = []
    for r in range(7, 18):
        raw = str(ws.cell(r, 5).value or "").strip()
        if not raw:
            empty_rows.append(r)
            continue
        for code in raw.split("/"):
            assert code.strip() in valid, f"R{r} 引用了未知底稿 {code!r}"
    assert empty_rows == [13, 17], f"源模板未填索引号的行应为 R13/R17，实得 {empty_rows}"


# ─── H0-7 舞弊迹象 19 条（与前端预置常量逐字） ───────────────────────────────


def test_h0_7_nineteen_fraud_indicators(wb):
    ws = wb["函证程序舞弊风险评价表H0-7"]
    items = [str(ws.cell(r, 1).value or "").strip() for r in range(6, 25)]
    assert len(items) == 19
    for i, text in enumerate(items, start=1):
        assert text.startswith(f"{i}."), f"第 {i} 条编号不符: {text[:20]!r}"
    assert str(ws.cell(25, 1).value or "").strip() == "……"
    assert str(ws["G5"].value).strip() == "是否存在"
    assert str(ws["H5"].value).strip() == "索引号或信息来源"
    assert str(ws["I5"].value).strip() == "应对措施"
    assert str(ws["H26"].value).strip() == "B50", "汇总行去向应为 B50"


def test_h0_7_matches_frontend_preset(wb):
    """19 条与 `fraudRiskPresets.PRESET_FRAUD_ITEMS` 逐字一致（去序号去尾标点）."""
    ws = wb["函证程序舞弊风险评价表H0-7"]
    src = strip_ts_comments(FRAUD_TS.read_text(encoding="utf-8"))
    block = ts_block(src, "PRESET_FRAUD_ITEMS: FraudRiskItem[] = [", "[", "]")
    fe = re.findall(r"description:\s*'([^']*)'", block)
    assert len(fe) == 19, f"前端预置 {len(fe)} 条"
    for i in range(19):
        xlsx = str(ws.cell(6 + i, 1).value or "").strip()
        xlsx = re.sub(r"^\d+\.", "", xlsx).rstrip("；;。")
        assert fe[i] == xlsx, f"第 {i + 1} 条不符:\n  前端={fe[i]!r}\n  源模板={xlsx!r}"


def test_h0_7_tooltip_rows_have_no_offset(wb):
    """🔴 H0-7 的两条举例挂在 **J19/J20**（与条目**同行**，无排版偏移）.

    D0-8/F0-8 是 J18/J19（比条目上移一行）→ `fraudRiskPresets` 的
    `item_14_reply_rate` / `item_15_independence` 映射在两种排版下都指向第 14/15 条，
    但依据不同。此处登记 H0-7 的事实，防后续按 D0-8 的偏移假设改错映射。
    """
    ws = wb["函证程序舞弊风险评价表H0-7"]
    assert str(ws["A19"].value or "").startswith("14."), "R19 应为第 14 条"
    assert "回函率" in str(ws["J19"].value or ""), "J19 应是回函率举例（与第 14 条同行）"
    assert str(ws["A20"].value or "").startswith("15."), "R20 应为第 15 条"
    assert "独立性" in str(ws["J20"].value or "") or "重大影响" in str(ws["J20"].value or "")
    # 反向自检：J18 无内容（证明确实没有偏移）
    assert not str(ws["J18"].value or "").strip(), "J18 应为空（H0-7 无偏移）"


# ─── H0-4 / H0-6 / H0-3 / H0-5 结构事实 ─────────────────────────────────────


def test_h0_4_nine_columns_and_total_row(wb):
    ws = wb["差异核对表H0-4"]
    assert [str(ws.cell(5, c).value or "").strip() for c in range(1, 10)] == [
        "询证函索引号", "被询证单位名称", "账户/交易", "发函金额/条款",
        "回函金额/条款", "差异（金额/条款）", "差异原因", "相关支持性证据", "是否调整",
    ]
    assert str(ws["A21"].value).strip() == "合计"


def test_h0_4_total_row_formulas(wb_formula):
    ws = wb_formula["差异核对表H0-4"]
    assert ws["D21"].value == "=SUM(D6:D20)"
    assert ws["E21"].value == "=SUM(E6:E20)"
    assert ws["F21"].value == "=SUM(F6:F20)"


def test_h0_6_fourteen_columns_with_group(wb):
    ws = wb["邮件传真回函可靠性验证H0-6"]
    assert str(ws["G5"].value).strip() == "期末未收回原件函证可靠性验证"
    assert "G5:M5" in {str(r) for r in ws.merged_cells.ranges}
    leaf = [str(ws.cell(6, c).value or "").strip() for c in range(7, 14)]
    assert leaf == [
        "被函证者身份确认（注1）", "发函及回函传真信息及验证", "发函邮箱", "回函邮箱",
        "邮箱可靠性验证（注2）", "是否致电被函证者确认", "对函证信息可靠性的考虑（注3）",
    ]
    assert str(ws["N5"].value).strip() == "回函可靠性结论"


def test_h0_2_second_send_block(wb):
    """H0-2 第二次发函信息 6 列 + 结果列（R8.1 的事实基线）."""
    ws = wb["核实被函证单位信息H0-2"]
    assert str(ws["AF5"].value).strip() == "第二次发函的被函证单位信息"
    assert "AF5:AK5" in {str(r) for r in ws.merged_cells.ranges}
    assert [str(ws.cell(6, c).value or "").strip() for c in range(32, 38)] == [
        "地址", "邮编", "联系人", "联系电话", "传真", "信息是否核查一致",
    ]
    assert "第二次发函结果" in str(ws["AL5"].value)
    # 被审计单位提供侧确有邮编 / 邮箱传真（R8.2）
    assert str(ws["E6"].value).strip() == "邮编"
    assert str(ws["H6"].value).strip() == "邮箱/传真"


def test_h0_3_three_checkpoints_and_signature(wb):
    ws = wb["跟函函证过程控制H0-3"]
    assert [str(ws.cell(r, 1).value or "").strip() for r in (23, 24, 25)] == [
        "是否了解处理函证的通常流程和处理人员",
        "是否确认询证函处理人员的身份及权限",
        "处理人员是否按正常流程处理",
    ]
    assert "审计项目组成员签名" in str(ws["A27"].value)


def test_h0_3_memo_mentions_escort_and_staff_no(wb):
    """跟函话术含**陪同情况**与**工号**（R9.1/R9.2 的事实基线）."""
    ws = wb["跟函函证过程控制H0-3"]
    body = "\n".join(str(ws.cell(r, 1).value or "") for r in range(10, 22))
    assert "无被审计单位人员陪同" in body
    assert "工号" in body
    assert "对外公开电话" in body, "应含第三方致电回访段（A18/A19）"


def test_h0_5_check_record_area_is_blank(wb):
    """🔴 H0-5「二、检查过程记录」R10:R19 是**空白自由区** —— 四区块属源外增强（R4.3）."""
    ws = wb["替代程序H0-5"]
    assert str(ws["A10"].value).strip() == "二、检查过程记录："
    for r in range(11, 20):
        row_vals = [ws.cell(r, c).value for c in range(1, ws.max_column + 1)]
        assert all(v in (None, "") for v in row_vals), f"H0-5 R{r} 竟有内容: {row_vals}"
    assert str(ws["A20"].value).strip() == "三、审计说明："
    assert str(ws["A23"].value).strip() == "四、审计结论："


def test_h0_5_sample_fields_and_guidance(wb):
    """H0-5 样本选取 6 字段 + 3 条替代程序编制说明（R4.2/R4.6 的事实基线）."""
    ws = wb["替代程序H0-5"]
    assert str(ws["A5"].value).strip() == "被函证单位名称："
    assert str(ws["A6"].value).strip() == "一、样本选取标准与规模："
    for anchor, label in [
        ("A7", "测试范围"), ("A8", "抽样总体"), ("A9", "抽样方法"),
        ("I7", "特定样本"), ("I8", "确定的抽样样本量"), ("I9", "抽样过程"),
    ]:
        assert str(ws[anchor].value or "").strip().rstrip("：") == label, anchor
    # 🔴 B7 是「借方发生额」单侧，**不含** F0-5 的「贷方发生额」措辞
    assert "借方发生额" in str(ws["B7"].value)
    assert "贷方发生额" not in str(ws["B7"].value), "H0-5 源模板 B7 无「贷方发生额」（那是 F0-5 措辞）"
    tips = [str(ws.cell(r, 2).value or "").strip() for r in (30, 31, 32)]
    assert tips == [
        "①检查本期付款、期后收货或回收；",
        "②检查原始凭证：合同、订货单、发票或收据、银行回单、支票存根等；",
        "③对回函可能性不高的、余额重大的，发函同时执行替代程序。",
    ]


# ─── 公式预设纠偏依据（R11.1） ──────────────────────────────────────────────


def test_prefill_preset_h0_sheet_points_to_real_tab():
    """H0 公式预设的 `sheet` 必须指向源模板真实存在的 tab（R11.1）.

    Task 21（`scripts/fix/fix_h0_prefill_presets.py`）已把 `审定表H0-1`（源模板无此 tab）
    改为 `函证结果汇总表H0-1`，故本条已转为正向断言（原 `xfail(strict=True)` 标记已移除
    —— 留着会 XPASS 报错，那正是设计意图）。
    """
    path = _REPO_ROOT / "backend" / "data" / "prefill_formula_mapping.json"
    data = json.loads(path.read_text(encoding="utf-8"))
    h0_blocks = [m for m in data["mappings"] if str(m.get("wp_code", "")) == "H0"]
    assert h0_blocks, "prefill_formula_mapping.json 无 H0 块"
    wb_names = set(openpyxl.load_workbook(H0_XLSX).sheetnames)
    for blk in h0_blocks:
        sheet = str(blk.get("sheet") or "")
        assert sheet, "H0 块缺 sheet 字段"
        # 只允许指向源模板真实存在的 tab
        assert sheet in wb_names, (
            f"H0 预设 sheet=「{sheet}」在源模板中不存在（真实汇总表名为『函证结果汇总表H0-1』）"
        )

"""CCR 源「改写」两项证据的现算（acnr-foundation-defect-closure T22，design §10.3 R2-F3）。

改写 = 把不可解析的 CCR 源的 source_sheet 改成目录里的某张表。每条改写都须两项证据：
- 证据①：模板册（xlsx / docx）里存在目标 tab —— 目录表的 sheet_name 或别名恰为该底稿某本模板册的 tab 名；
- 证据②：CCR 的字段文字出现在该 tab 的文本里。字段文字取 source_cell_label；没有标签时取 source_cell；
  形如 A1 / 程序行号（R38）的不算字段文字（它在任何表里都可能出现，证明不了是这张表）。

模板册按文件名以底稿编码开头（后面不紧跟字母数字）认定属于该底稿。测试与豁免清单的证据都由本模块现算。
"""
from __future__ import annotations

import re
import warnings
from functools import lru_cache
from pathlib import Path
from typing import Any

_BACKEND = Path(__file__).resolve().parents[2]
TEMPLATES = _BACKEND / "wp_templates"
_A1 = re.compile(r"^[A-Z]{1,3}[0-9]{1,7}$")


def field_text(ref: dict[str, Any]) -> str | None:
    for raw in (ref.get("source_cell_label"), ref.get("source_cell")):
        text = str(raw or "").strip()
        if text and not _A1.match(text.upper()):
            return text
    return None


@lru_cache(maxsize=1)
def _template_files() -> tuple[Path, ...]:
    return tuple(sorted(p for p in TEMPLATES.rglob("*") if p.suffix.lower() in (".xlsx", ".xlsm", ".docx")
                        and not p.name.startswith("~$")))


def files_of(wp_code: str) -> list[Path]:
    pat = re.compile(rf"^{re.escape(wp_code)}(?![A-Za-z0-9])")
    return [p for p in _template_files() if pat.match(p.name)]


@lru_cache(maxsize=None)
def tabs_text(path: Path) -> dict[str, str]:
    """模板册 → {tab 名: 文本}；docx 整本记在键「<docx>」下。"""
    out: dict[str, str] = {}
    with warnings.catch_warnings():
        warnings.simplefilter("ignore")
        if path.suffix.lower() == ".docx":
            import docx

            d = docx.Document(str(path))
            parts = [p.text for p in d.paragraphs]
            parts += [c.text for t in d.tables for row in t.rows for c in row.cells]
            out["<docx>"] = "\n".join(parts)
        else:
            import openpyxl

            wb = openpyxl.load_workbook(str(path), read_only=True, data_only=False)
            try:
                for ws in wb.worksheets:
                    out[ws.title] = "\n".join(v for row in ws.iter_rows(values_only=True) for v in row if isinstance(v, str))
            finally:
                wb.close()
    return out


def _rel(p: Path) -> str:
    return p.relative_to(_BACKEND).as_posix()


def target_evidence(ref: dict[str, Any], sheet: dict[str, Any]) -> dict[str, Any]:
    """某张目录表作为改写目标时的两项证据。"""
    names = {sheet.get("sheet_name"), *(sheet.get("sheet_name_aliases") or [])}
    text = field_text(ref)
    tabs = [(p, tab, body) for p in files_of(str(ref.get("source_wp") or "")) for tab, body in tabs_text(p).items()
            if tab in names]
    return {"tab_found": bool(tabs), "text_found": bool(text and any(text in body for _p, _t, body in tabs)),
            "files": sorted({_rel(p) for p, _t, _b in tabs})}


def source_evidence(ref: dict[str, Any], sheets_of_parent: list[dict[str, Any]]) -> dict[str, Any]:
    """不可解析的 CCR 源：逐个目录候选表现算两项证据，归纳缺哪一项。"""
    wp, sheet = str(ref.get("source_wp") or ""), str(ref.get("source_sheet") or "")
    text = field_text(ref)
    per = {s["addr_id"]: target_evidence(ref, s) for s in sorted(sheets_of_parent, key=lambda x: x["addr_id"])}
    tab_found = sorted(a for a, e in per.items() if e["tab_found"])
    text_found = sorted(a for a, e in per.items() if e["text_found"])
    source_tab_files = sorted(_rel(p) for p in files_of(wp) if sheet in tabs_text(p))
    if text_found:
        missing = "none"
    elif not tab_found:
        missing = "evidence_1"
    elif text is None:
        missing = "evidence_2_no_field_text"
    else:
        missing = "evidence_2"
    return {"field_text": text, "candidate_sheets": len(per), "tab_found": tab_found, "text_found": text_found,
            "source_tab_files": source_tab_files, "missing": missing}


__all__ = ["field_text", "files_of", "source_evidence", "tabs_text", "target_evidence"]

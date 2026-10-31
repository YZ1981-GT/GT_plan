"""目录单元格坐标复核守卫（spec wp-deeplink-locate-coverage-closure 需求 1.5，design §2.3；
spec wp-deeplink-locate-residual-closure 需求 4.2 / 5.2，design §4）。

1. ``review_cell_coordinates.py --check``：停用判定（kind=seed_coordinate_mismatch）与 ``t_address_registry_seed.json``
   都等于读模板现算（模板、旧种子或改指规则变了即红）；
2. 目录里未停用的单元格全部来自 t_ 种子（以后新增坐标不经模板复核即红）；T06 的「语义待核对」判定已由复核接管，一条不留；
3. t_ 种子每个坐标按其 rule 成立：``conclusion_label`` 的模板现值是「审计结论」标签；``explicit`` 只来自显式改指表；
   ``total_row`` 用 openpyxl 独立复算（唯一合计行、表头恰等于指标的唯一列、目标格是公式或数值）；
4. 本 kind 的停用格都在目录、已停用、``replaced_by`` 是父 sheet 或 t_ 种子里声明替代它的新坐标（两边互相对得上）；
5. （pg_only）PG 公式类列对本 kind 停用格的 addr_id / formula_ref 引用为 0（与 T05 孤儿判定同口径）；
6. 前端 ``generated/catalogTextAnchors.ts`` 与目录现算一致（``generate_cell_text_anchors.py --check``），只含复核过的
   「审计结论」格；该 ``--check`` 接在 CI 作业 ``acnr-cell-text-anchors-generated`` 上（需求 2.2）。
"""
from __future__ import annotations

import json
import re
from functools import lru_cache
from pathlib import Path

import pytest

from app.services.acnr.catalog import get_catalog
from scripts.acnr import review_cell_coordinates as rcc

_BACKEND = Path(__file__).resolve().parents[2]
_ACNR = _BACKEND / "data" / "acnr"


@lru_cache(maxsize=1)
def _verdicts() -> list[dict]:
    raw = json.loads((_ACNR / "sources" / "sheet_verdicts.json").read_text(encoding="utf-8"))
    return raw["cell_verdicts"]


@lru_cache(maxsize=1)
def _seed() -> dict:
    return json.loads(rcc.T_SEED.read_text(encoding="utf-8"))


def _seed_coords() -> dict[str, tuple[dict, dict]]:
    """t_ 种子坐标 addr_id → (entry, coordinate)。"""
    out = {}
    for e in _seed()["entries"]:
        for c in e["coordinates"]:
            out[f"{rcc._parent_code(e['wp_code'])}/{e['wp_code']}/{c['cell_address']}"] = (e, c)
    return out


def test_check_mode_matches_templates():
    assert rcc.run(write=False) == 0


def test_live_cells_come_from_reviewed_seed():
    cat = get_catalog()
    live = {a for a, c in cat.cells_by_addr_id.items() if not c.get("deprecated")}
    stray = sorted(live - set(_seed_coords()))
    assert stray == [], f"未经模板复核的有效单元格（须经 review_cell_coordinates.py 登记）：{stray[:10]}"
    assert set(_seed_coords()) <= live, "t_ 种子坐标应全部是目录里的有效单元格"
    assert len(_seed_coords()) >= 3 and {"E1/E1-14/I18", "E1/E1-14/C34", "E1/E1-14/A50"} <= set(_seed_coords())


def test_no_semantic_unverified_verdict_left():
    """T06 原判「语义待核对」的 33 格已纳入复核（同一格不留两份判定），且都在本 kind 里有判定。"""
    keep = [v["addr_id"] for v in _verdicts() if v["verdict"] == rcc.SUPERSEDED_VERDICT]
    assert keep == [], keep[:5]
    mine = {v["addr_id"] for v in _verdicts() if v.get("kind") == rcc.KIND}
    assert {"G8/G8-4/B35", "K1/K1-2/I30", "S15/S15/C12"} <= mine


_TOTAL_LABELS = {"合计", "总计", "合计数"}


def _nospace(v) -> str:
    return re.sub(r"\s+", "", str(v or ""))


def _total_row_holds(c: dict, seed_labels: list[str]) -> list[str]:
    """total_row 坐标的独立复算（不调脚本的规则函数）：返回不成立的原因。"""
    import openpyxl
    from openpyxl.utils.cell import column_index_from_string, coordinate_from_string

    ev = c["evidence"]
    ws = openpyxl.load_workbook(_BACKEND.parent / ev["template_file"], data_only=False)[ev["template_tab"]]
    merged = list(ws.merged_cells.ranges)

    def val(r: int, col: int):
        m = next((m for m in merged if m.min_row <= r <= m.max_row and m.min_col <= col <= m.max_col), None)
        return ws.cell(row=m.min_row, column=m.min_col).value if m else ws.cell(row=r, column=col).value

    def first_text(r: int) -> str:
        for col in range(1, (ws.max_column or 0) + 1):
            v = val(r, col)
            if isinstance(v, str) and v.strip():
                return _nospace(v)
        return ""

    col_letter, row = coordinate_from_string(c["cell_address"])
    col = column_index_from_string(col_letter)
    metric = re.sub(r"合计$", "", c["description"])
    problems = []
    totals = [r for r in range(1, (ws.max_row or 0) + 1) if first_text(r) in _TOTAL_LABELS]
    if totals != [row]:
        problems.append(f"合计行不唯一或不是本行：{totals}")
    if _nospace(val(ev["header_row"], col)) != metric or not (row - 15 <= ev["header_row"] < row):
        problems.append(f"表头 {ev['header_row']} 行不是「{metric}」或不在合计行以上 15 行内")
    cols = {cc for cc in range(1, (ws.max_column or 0) + 1)
            if any(_nospace(val(rr, cc)) == metric for rr in range(max(1, row - 15), row))}
    if cols != {col}:
        problems.append(f"表头等于「{metric}」的列不唯一：{sorted(cols)}")
    v = ws.cell(row=row, column=col).value
    if isinstance(v, bool) or not (isinstance(v, (int, float)) or (isinstance(v, str) and v.startswith("="))):
        problems.append(f"目标格不是公式或数值：{v!r}")
    for label in seed_labels:
        s = _nospace(label)
        if "-" in s and (s.split("-", 1)[0] == "合计行" or s.split("-", 1)[0].endswith("合计")):
            s = s.split("-", 1)[1]
        if re.sub(r"合计$", "", s) != metric:
            problems.append(f"所替代种子的语义「{label}」与指标「{metric}」不符")
    return problems


def test_seed_coordinates_satisfy_their_rule():
    explicit = {(old.split("/")[1], v["a1"]): old for old, v in rcc.EXPLICIT.items()}
    seed_label = {v["addr_id"]: (v.get("evidence") or {}).get("seed_label") or "" for v in _verdicts()}
    rules: dict[str, int] = {}
    for aid, (entry, c) in _seed_coords().items():
        rules[c["rule"]] = rules.get(c["rule"], 0) + 1
        if c["rule"] == "conclusion_label":
            text = re.sub(r"\s+", "", str(c["evidence"].get("cell_value") or ""))
            assert rcc.CONCLUSION_RE.match(text), (aid, text)
            assert (c["description"], c["purpose"]) == (rcc.CONCLUSION_LABEL, "conclusion"), aid
        elif c["rule"] == "total_row":
            problems = _total_row_holds(c, [seed_label[old] for old in c["replaces"]])
            assert problems == [], (aid, problems)
        else:
            assert c["rule"] == "explicit" and (entry["wp_code"], c["cell_address"]) in explicit, aid
            assert c["replaces"] == [explicit[(entry["wp_code"], c["cell_address"])]], aid
        assert c["replaces"] and c["evidence"]["template_tab"] == entry["sheet_name"], aid
    # 防空转：三种规则都有（合计行规则现算 D1-3 J21、D2-2 Q26）
    assert set(rules) == {"conclusion_label", "explicit", "total_row"} and rules["total_row"] >= 2, rules


def _sheet(rows: list[list]):
    import openpyxl

    ws = openpyxl.Workbook().active
    for r, vals in enumerate(rows, start=1):
        for c, v in enumerate(vals, start=1):
            if v is not None:
                ws.cell(row=r, column=c, value=v)
    return ws


def test_total_row_rule_conditions_on_synthetic_sheets():
    """合计行规则（design §4.2）每个条件都要成立才改指：唯一合计行、合计行以上 15 行内表头恰等于指标的唯一列、
    目标格是公式或数值；种子语义不含「合计」不算合计类。"""
    head = ["项目", "期初余额", "期末余额"]
    body = [["甲", 1, 2], ["乙", 3, 4]]
    ok = _sheet([head, *body, ["合计", "=SUM(B2:B3)", "=SUM(C2:C3)"]])
    assert rcc.total_target(ok, "合计行-期末余额合计") == ("C4", 1)
    assert rcc.total_target(ok, "期初余额合计") == ("B4", 1)
    assert rcc.total_target(ok, "银行存款合计-期末余额") == ("C4", 1)
    assert rcc.total_target(_sheet([head, *body, [" 合 计 ", 5, 6]]), "合计行-期末余额合计") == ("C4", 1)
    assert rcc.total_target(ok, "合计行-本期增加合计") is None, "没有该表头"
    assert rcc.total_target(ok, "期末余额") is None, "语义不含「合计」不算合计类"
    assert rcc.total_target(_sheet([head, *body, ["合计", 1, 2], ["小计"], ["总计", 1, 2]]), "合计行-期末余额合计") is None, \
        "合计行不唯一"
    assert rcc.total_target(_sheet([["项目", "期末余额", "期末余额"], *body, ["合计", 1, 2]]), "合计行-期末余额合计") is None, \
        "表头等于指标的列不唯一"
    assert rcc.total_target(_sheet([head, *body, ["合计", "见附表", "—"]]), "合计行-期末余额合计") is None, "目标格是文字"
    assert rcc.total_target(_sheet([head, *body, ["合计", True, True]]), "合计行-期末余额合计") is None, "目标格是布尔"
    assert rcc.total_target(_sheet([head, *[["行", 1, 1]] * 16, ["合计", 1, 2]]), "合计行-期末余额合计") is None, \
        "表头在合计行以上 15 行之外"


def test_merged_verdicts_drop_superseded_semantic_unverified():
    """写出时：本 kind 复核到的格，别的 kind 里的「语义待核对」判定删掉（同一格不留两份）；复核没覆盖到的不删；其余判定不动。"""
    raw = {"_doc": "", "cell_verdicts": [
        {"addr_id": "X1/X1-1/A1", "kind": "seed_name_mismatch", "verdict": rcc.SUPERSEDED_VERDICT},
        {"addr_id": "X1/X1-1/A2", "kind": "seed_name_mismatch", "verdict": rcc.SUPERSEDED_VERDICT},
        {"addr_id": "X1/X1-1/B9", "kind": "seed_name_mismatch", "verdict": "deprecate"},
        {"addr_id": "X1/X1-1/C3", "kind": "orphan", "verdict": "deprecate"},
        {"addr_id": "X1/X1-1/Z1", "kind": rcc.KIND, "verdict": "deprecate"},
    ]}
    mine = [{"addr_id": "X1/X1-1/A1", "kind": rcc.KIND, "verdict": "deprecate"}]
    out = rcc._merged_verdicts(raw, mine)
    assert [(v["addr_id"], v["kind"]) for v in out["cell_verdicts"]] == [
        ("X1/X1-1/A2", "seed_name_mismatch"), ("X1/X1-1/B9", "seed_name_mismatch"), ("X1/X1-1/C3", "orphan"),
        ("X1/X1-1/A1", rcc.KIND)]
    assert rcc.KIND in out["_doc"] and rcc._DOC_NOTE_SUPERSEDE in out["_doc"]
    assert [v["addr_id"] for v in rcc._superseded(raw, mine)] == ["X1/X1-1/A1"]


def test_review_scope_takes_over_semantic_unverified_cells():
    """复核范围：别的 kind 已判定的不复核，「语义待核对」的除外；本 kind 自己的与没判定的都复核。"""
    cells = [{"addr_id": f"X1/X1-1/A{i}"} for i in range(1, 5)]
    raw = {"cell_verdicts": [
        {"addr_id": "X1/X1-1/A1", "kind": "orphan", "verdict": "deprecate"},
        {"addr_id": "X1/X1-1/A2", "kind": "seed_name_mismatch", "verdict": rcc.SUPERSEDED_VERDICT},
        {"addr_id": "X1/X1-1/A3", "kind": rcc.KIND, "verdict": "deprecate"},
    ]}
    assert [c["addr_id"] for c in rcc.review_scope(raw, cells)] == ["X1/X1-1/A2", "X1/X1-1/A3", "X1/X1-1/A4"]


def test_check_lines_name_leftover_semantic_unverified_verdict():
    """--check：被取代的「语义待核对」判定还留着即漂移并点名；_doc 缺说明也点名。"""
    mine = [{"addr_id": "X1/X1-1/A2", "kind": rcc.KIND, "verdict": "deprecate"}]
    seed = {"entries": []}
    doc = rcc._DOC_NOTE + rcc._DOC_NOTE_SUPERSEDE
    assert rcc.check_lines({"_doc": doc, "cell_verdicts": list(mine)}, mine, seed, seed) == []
    left = {"_doc": doc, "cell_verdicts": [
        *mine, {"addr_id": "X1/X1-1/A2", "kind": "seed_name_mismatch", "verdict": rcc.SUPERSEDED_VERDICT}]}
    assert [ln for ln in rcc.check_lines(left, mine, seed, seed) if rcc.SUPERSEDED_VERDICT in ln and "X1/X1-1/A2" in ln]
    no_note = rcc.check_lines({"_doc": rcc._DOC_NOTE, "cell_verdicts": list(mine)}, mine, seed, seed)
    assert no_note == ["sheet_verdicts._doc: 缺「语义待核对」判定由本 kind 取代的说明"]


def test_deprecated_cells_point_to_parent_or_declared_replacement():
    cat = get_catalog()
    coords = _seed_coords()
    declared = {old: new for new, (_e, c) in coords.items() for old in c["replaces"]}
    mine = [v for v in _verdicts() if v.get("kind") == rcc.KIND]
    assert len(mine) >= 300, "防空转：本 kind 的停用判定应有数百条"
    bad = []
    for v in mine:
        cell = cat.cells_by_addr_id.get(v["addr_id"]) or {}
        want = declared.get(v["addr_id"], v["parent_sheet_code"])
        if not (cell.get("deprecated") and cell.get("replaced_by") == v["replaced_by"] == want):
            bad.append((v["addr_id"], v["replaced_by"], want, cell.get("deprecated")))
    assert bad == [], bad[:5]
    assert set(declared) <= {v["addr_id"] for v in mine}, "t_ 种子声明替代的旧坐标都应有停用判定"


def test_text_anchor_artifact_matches_catalog():
    from scripts.acnr import generate_cell_text_anchors as gta

    assert gta.main(["--check"]) == 0, "catalogTextAnchors.ts 与目录现算不一致：运行 generate_cell_text_anchors.py --write"
    anchors = gta.build(json.loads(gta.CATALOG.read_text(encoding="utf-8")))
    ids = {f"{rcc._parent_code(code)}/{code}/{a1}" for code, cells in anchors.items() for a1 in cells}
    assert len(ids) >= 60 and "E1/E1-14/A50" in ids, "防空转：结论格文字锚点应有数十条且含 E1-14 A50"
    assert {t for cells in anchors.values() for t in cells.values()} == {"审计结论"}
    conclusion_seed = {a for a, (_e, c) in _seed_coords().items() if c["rule"] == "conclusion_label"}
    assert ids <= conclusion_seed, sorted(ids - conclusion_seed)[:5]
    assert "G8/G8-4/A24" in ids, "G8-4 B35 改指的「审计结论」标签格应进文字锚点对照"


def test_text_anchor_check_wired_into_ci():
    import yaml

    wf = _BACKEND.parent / ".github" / "workflows" / "governance-checks.yml"
    job = (yaml.safe_load(wf.read_bytes().decode("utf-8")).get("jobs") or {}).get("acnr-cell-text-anchors-generated") or {}
    assert job, "governance-checks.yml 缺作业 acnr-cell-text-anchors-generated"
    assert not job.get("continue-on-error")
    assert any(s.get("run") == "python backend/scripts/acnr/generate_cell_text_anchors.py --check"
               for s in job.get("steps") or [])


def _pg_texts() -> list[str]:
    import psycopg2

    conn = psycopg2.connect(dbname="audit_platform", user="postgres", password="postgres",
                            host="localhost", port=5432, connect_timeout=5)
    conn.set_session(readonly=True, autocommit=True)
    try:
        cur = conn.cursor()
        cur.execute("SELECT table_name, column_name FROM information_schema.columns WHERE table_schema='public' "
                    "AND (column_name ILIKE '%%formula%%' OR column_name ILIKE '%%expression%%') "
                    "AND data_type IN ('text','character varying','jsonb','json')")
        cols = cur.fetchall()
        assert cols, "没有找到公式类列（防空转）"
        texts = []
        for t, c in cols:
            cur.execute(f'SELECT "{c}"::text FROM "{t}" WHERE "{c}" IS NOT NULL')
            texts += [r[0] for r in cur.fetchall() if r[0]]
        return texts
    finally:
        conn.close()


@pytest.mark.pg_only
def test_no_pg_formula_references_to_deprecated_seed_coordinates():
    cat = get_catalog()
    texts = _pg_texts()
    hits = {}
    for v in _verdicts():
        if v.get("kind") != rcc.KIND:
            continue
        fr = (cat.cells_by_addr_id.get(v["addr_id"]) or {}).get("formula_ref") or ""
        n = sum(1 for t in texts for needle in (v["addr_id"], fr) if needle and needle in t)
        if n:
            hits[v["addr_id"]] = n
    assert hits == {}, f"停用的种子坐标在 PG 公式里有引用（不得静默停用）：{hits}"

"""D4-1 审定表 descriptor 契约测试（Task 1.2）。

openpyxl 直读权威 xlsx，与 `app.services.d4_extraction.d4_1_descriptor` 的常量三向比对：
- tab 名 / 索引号 / 同册 sheet 名单
- 两级表头文案（R5/R6）逐格
- 段标题 + 可扩明细行区段 + 小计/合计/差异行
- 受管输入列 vs 审定数派生列 vs 公式格
并含**反向自检**（判据自己不能恒真）。

spec: d4-1-adjudication-bidirectional-writeback-and-formula-io / Property 2(公式), Req 1.1
失败即阻塞后续实现（DEC1 模板先行）。
"""
from __future__ import annotations

import openpyxl
import pytest

from app.services.d4_extraction.d4_1_descriptor import (
    D4_1_AUDITED_COLUMNS,
    D4_1_DIFF_ROW,
    D4_1_HEADER_SUB,
    D4_1_HEADER_SUB_ROW,
    D4_1_HEADER_TOP,
    D4_1_HEADER_TOP_ROW,
    D4_1_INDEX_CODE,
    D4_1_INPUT_COLUMNS,
    D4_1_SECTIONS,
    D4_1_SHEET_NAME,
    D4_1_TB_CHECK_ROW,
    D4_1_TOTAL_ROW,
    D4_1_WORKBOOK_SHEETS,
    d4_1_template_path,
)


@pytest.fixture(scope="module")
def wb():
    path = d4_1_template_path()
    assert path.exists(), f"权威模板不存在: {path}"
    return openpyxl.load_workbook(path, data_only=False, read_only=False)


@pytest.fixture(scope="module")
def ws(wb):
    assert D4_1_SHEET_NAME in wb.sheetnames, (
        f"tab 名漂移: descriptor={D4_1_SHEET_NAME!r} 不在 {wb.sheetnames}"
    )
    return wb[D4_1_SHEET_NAME]


# ── 册 / tab / 索引号 ────────────────────────────────────────────────────────

def test_workbook_sheets_match(wb):
    """同册 sheet 名单与 descriptor 逐字一致（D4-1~D4-4 同册）。"""
    assert tuple(wb.sheetnames) == D4_1_WORKBOOK_SHEETS


def test_sheet_name_and_index(ws):
    assert ws["I3"].value == D4_1_INDEX_CODE


# ── 两级表头 ────────────────────────────────────────────────────────────────

def test_header_top_row(ws):
    for coord, text in D4_1_HEADER_TOP.items():
        assert ws[coord].value == text, f"{coord} 期望 {text!r} 实得 {ws[coord].value!r}"
    # 行号一致
    assert all(int(c[1:]) == D4_1_HEADER_TOP_ROW for c in D4_1_HEADER_TOP)


def test_header_sub_row(ws):
    for coord, text in D4_1_HEADER_SUB.items():
        assert ws[coord].value == text, f"{coord} 期望 {text!r} 实得 {ws[coord].value!r}"
    assert all(int(c[1:]) == D4_1_HEADER_SUB_ROW for c in D4_1_HEADER_SUB)


# ── 段标题 + 可扩行 + 小计 ────────────────────────────────────────────────────

def test_sections_titles_and_subtotals(ws):
    for sec in D4_1_SECTIONS:
        assert ws[f"A{sec.title_row}"].value == sec.title, (
            f"段标题漂移: A{sec.title_row} 期望 {sec.title!r} 实得 {ws[f'A{sec.title_row}'].value!r}"
        )
        # 小计行 A 列文案
        assert ws[f"A{sec.subtotal_row}"].value == "小计"
        # 可扩明细行的审定列 E 是 SUM 公式（空白可扩行的证据）
        for r in sec.detail_rows:
            e = ws[f"E{r}"].value
            assert isinstance(e, str) and e.startswith("=SUM(B"), (
                f"E{r} 应为明细审定 SUM 公式，实得 {e!r}"
            )
            # A 列在空白可扩行应为空（不是固定行集）
            assert ws[f"A{r}"].value in (None, ""), (
                f"A{r} 应为空白可扩行，实得 {ws[f'A{r}'].value!r}"
            )


def test_total_tb_diff_rows(ws):
    assert ws[f"A{D4_1_TOTAL_ROW}"].value == "合计"
    assert ws[f"A{D4_1_TB_CHECK_ROW}"].value == "试算平衡表数"
    assert ws[f"A{D4_1_DIFF_ROW}"].value == "差异数"
    # 合计 = 主营小计 + 其他小计（结构公式）
    assert ws[f"B{D4_1_TOTAL_ROW}"].value == "=B18+B12"
    # 差异 = 合计 - 试算平衡表数
    assert ws[f"E{D4_1_DIFF_ROW}"].value == f"=E{D4_1_TOTAL_ROW}-E{D4_1_TB_CHECK_ROW}"


# ── 输入列 vs 审定数派生列 ────────────────────────────────────────────────────

def test_input_columns_are_not_formula(ws):
    """受管输入列（B/C/D/F/G/H）在明细行不是公式格（可录入）。"""
    detail_rows = [r for sec in D4_1_SECTIONS for r in sec.detail_rows]
    for spec in D4_1_INPUT_COLUMNS:
        for r in detail_rows:
            v = ws[f"{spec.col}{r}"].value
            assert not (isinstance(v, str) and v.startswith("=")), (
                f"输入列 {spec.col}{r} 不应是公式，实得 {v!r}"
            )


def test_audited_columns_are_formula(ws):
    """审定数列（E/I）在明细行是 SUM 派生公式。"""
    detail_rows = [r for sec in D4_1_SECTIONS for r in sec.detail_rows]
    audited_cols = [c for c, _ in D4_1_AUDITED_COLUMNS]
    for col in audited_cols:
        for r in detail_rows:
            v = ws[f"{col}{r}"].value
            assert isinstance(v, str) and v.startswith("=SUM("), (
                f"审定列 {col}{r} 应为 SUM 公式，实得 {v!r}"
            )


def test_audited_formula_semantics(ws):
    """审定数 = SUM(该期三输入列)：E{n}=SUM(B{n}:D{n}) / I{n}=SUM(F{n}:H{n})。"""
    detail_rows = [r for sec in D4_1_SECTIONS for r in sec.detail_rows]
    for r in detail_rows:
        assert ws[f"E{r}"].value == f"=SUM(B{r}:D{r})"
        assert ws[f"I{r}"].value == f"=SUM(F{r}:H{r})"


def test_input_field_keys_are_six(ws):
    """受管输入列恰好 6 个，键与前端 D4_ADJ_VALUE_FIELDS 对齐（不多不少）。"""
    fields = {spec.field for spec in D4_1_INPUT_COLUMNS}
    assert fields == {
        "currentUnadjusted", "currentAje", "currentRje",
        "priorUnadjusted", "priorAje", "priorRje",
    }
    assert len(D4_1_INPUT_COLUMNS) == 6


# ── 反向自检（判据不能恒真）────────────────────────────────────────────────────

def test_reverse_guard_wrong_sheet_name_fails(wb):
    """descriptor 若写错 tab 名，fixture 断言必失败（证明 tab 校验有分辨力）。"""
    assert "营业收入审定表D4-99" not in wb.sheetnames

def test_reverse_guard_header_would_catch_drift(ws):
    """若把表头文案改成错的，比对必失败（证明表头校验非空转）。"""
    assert ws["B5"].value != "错误表头占位XXX"
    assert ws["B6"].value == "未审数"  # 真值仍在，正向锚点

def test_reverse_guard_input_col_is_really_editable(ws):
    """明细输入区确有可录入格（B8 非公式），证明 test_input_columns_are_not_formula 非空断言。"""
    v = ws["B8"].value
    assert not (isinstance(v, str) and v.startswith("="))

"""Property-Based Tests for wp_docx_template_parser using hypothesis.

Property 1: Placeholder extraction completeness
Property 2: Parser structural completeness
Property 3: Parser mtime cache idempotence
Property 4: Placeholder extraction round-trip

**Validates: Requirements 2.1, 2.2, 2.3, 2.5, 2.6**
"""

from __future__ import annotations

import os
import tempfile
import time
from pathlib import Path

from docx import Document
from hypothesis import HealthCheck, given, settings
from hypothesis import strategies as st

from app.services.wp_docx_template_parser import (
    TemplateStructure,
    format_placeholder_summary,
    get_cached_structure,
    invalidate_cache,
    parse_template,
)

# ─── Hypothesis Strategies ────────────────────────────────────────────────────

# Valid field_id: starts with letter/underscore, alphanumeric + underscore
_field_id_st = st.from_regex(r"[a-z][a-z0-9_]{2,15}", fullmatch=True)

# Label: Chinese or ASCII text for display
_label_st = st.text(
    alphabet=st.characters(whitelist_categories=("L", "N"), whitelist_characters="，。"),
    min_size=2,
    max_size=10,
)


@st.composite
def placeholder_list_strategy(draw: st.DrawFn) -> list[tuple[str, str]]:
    """Generate a list of unique (field_id, label) pairs for placeholders."""
    count = draw(st.integers(min_value=1, max_value=5))
    ids = draw(
        st.lists(_field_id_st, min_size=count, max_size=count, unique=True)
    )
    labels = draw(st.lists(_label_st, min_size=count, max_size=count))
    return list(zip(ids, labels))


def _create_docx_with_placeholders(placeholders: list[tuple[str, str]]) -> str:
    """Create a temp docx file with the given placeholders. Returns file path."""
    doc = Document()
    doc.add_heading("测试文档标题", level=1)
    for field_id, label in placeholders:
        doc.add_paragraph(f"请填写 ${{{field_id}:{label}}} 的内容。")
    # Add a table with one placeholder
    if placeholders:
        table = doc.add_table(rows=2, cols=2)
        table.cell(0, 0).text = "字段"
        table.cell(0, 1).text = "值"
        table.cell(1, 0).text = placeholders[0][1]
        table.cell(1, 1).text = f"${{{placeholders[0][0]}:{placeholders[0][1]}}}"

    fd, path = tempfile.mkstemp(suffix=".docx")
    os.close(fd)
    doc.save(path)
    return path


# ─── Property 1: Placeholder extraction completeness ─────────────────────────


class TestProperty1PlaceholderExtractionCompleteness:
    """Property 1: Placeholder extraction completeness.

    For any generated docx containing ${field_id} placeholders,
    parse_template SHALL extract ALL of them.

    **Validates: Requirements 2.1, 2.3**
    """

    @settings(max_examples=5)
    @given(placeholders=placeholder_list_strategy())
    def test_all_placeholders_extracted(self, placeholders: list[tuple[str, str]]):
        """Every ${field_id:label} placeholder in the docx is extracted."""
        path = _create_docx_with_placeholders(placeholders)
        try:
            result = parse_template(path)
            extracted_ids = {p.field_id for p in result.placeholders}
            expected_ids = {fid for fid, _ in placeholders}
            assert expected_ids.issubset(extracted_ids), (
                f"Missing placeholders: {expected_ids - extracted_ids}. "
                f"Expected: {expected_ids}, Got: {extracted_ids}"
            )
        finally:
            os.unlink(path)

    @settings(max_examples=5)
    @given(placeholders=placeholder_list_strategy())
    def test_extracted_placeholders_have_valid_fields(self, placeholders: list[tuple[str, str]]):
        """Each extracted placeholder has field_id, label, data_type, default_value, position."""
        path = _create_docx_with_placeholders(placeholders)
        try:
            result = parse_template(path)
            for p in result.placeholders:
                assert p.field_id, "field_id must be non-empty"
                assert p.label, "label must be non-empty"
                assert p.data_type in ("text", "date", "textarea", "number"), (
                    f"Invalid data_type: {p.data_type}"
                )
                assert p.default_value, "default_value must be non-empty"
                assert isinstance(p.position, dict), "position must be a dict"
        finally:
            os.unlink(path)


# ─── Property 2: Parser structural completeness ──────────────────────────────


class TestProperty2StructuralCompleteness:
    """Property 2: Parser structural completeness.

    For any valid docx, the parsed TemplateStructure SHALL contain paragraphs
    with text/style/heading_level/placeholder_ids, tables with rows/placeholder_ids,
    and placeholders with field_id/label/data_type/default_value/position.

    **Validates: Requirements 2.2, 2.3**
    """

    @settings(max_examples=5)
    @given(placeholders=placeholder_list_strategy())
    def test_paragraphs_have_required_fields(self, placeholders: list[tuple[str, str]]):
        """Each paragraph has index, text, style, heading_level, placeholder_ids."""
        path = _create_docx_with_placeholders(placeholders)
        try:
            result = parse_template(path)
            assert len(result.paragraphs) > 0, "paragraphs list must not be empty"
            for para in result.paragraphs:
                assert isinstance(para.index, int), "index must be int"
                assert isinstance(para.text, str), "text must be str"
                assert isinstance(para.style, str), "style must be str"
                assert isinstance(para.heading_level, int), "heading_level must be int"
                assert 0 <= para.heading_level <= 6, f"heading_level out of range: {para.heading_level}"
                assert isinstance(para.placeholder_ids, list), "placeholder_ids must be list"
        finally:
            os.unlink(path)

    @settings(max_examples=5)
    @given(placeholders=placeholder_list_strategy())
    def test_tables_have_required_fields(self, placeholders: list[tuple[str, str]]):
        """Each table has index, rows, placeholder_ids."""
        path = _create_docx_with_placeholders(placeholders)
        try:
            result = parse_template(path)
            assert len(result.tables) > 0, "tables list must not be empty"
            for tbl in result.tables:
                assert isinstance(tbl.index, int), "table index must be int"
                assert isinstance(tbl.rows, list), "rows must be list"
                assert all(isinstance(row, list) for row in tbl.rows), "each row must be list"
                assert isinstance(tbl.placeholder_ids, list), "placeholder_ids must be list"
        finally:
            os.unlink(path)

    @settings(max_examples=5)
    @given(placeholders=placeholder_list_strategy())
    def test_result_is_template_structure(self, placeholders: list[tuple[str, str]]):
        """Result is a TemplateStructure with paragraphs, tables, placeholders, metadata."""
        path = _create_docx_with_placeholders(placeholders)
        try:
            result = parse_template(path)
            assert isinstance(result, TemplateStructure)
            assert isinstance(result.paragraphs, list)
            assert isinstance(result.tables, list)
            assert isinstance(result.placeholders, list)
            assert isinstance(result.metadata, dict)
            assert "template_name" in result.metadata
            assert "last_parsed_at" in result.metadata
        finally:
            os.unlink(path)


# ─── Property 3: Parser mtime cache idempotence ──────────────────────────────


class TestProperty3CacheIdempotence:
    """Property 3: Parser mtime cache idempotence.

    For any template file path, calling get_cached_structure twice without
    modifying the file SHALL return the SAME TemplateStructure object (identity).
    When the file mtime changes, the next call SHALL re-parse.

    **Validates: Requirements 2.5**
    """

    @settings(max_examples=5)
    @given(field_id=_field_id_st)
    def test_same_file_returns_same_object(self, field_id: str):
        """Two calls to get_cached_structure on same unmodified file return identical object."""
        invalidate_cache()

        doc = Document()
        doc.add_paragraph(f"${{{field_id}:label}}")

        fd, path = tempfile.mkstemp(suffix=".docx")
        os.close(fd)
        doc.save(path)

        try:
            result1 = get_cached_structure(path, "TEST-WP")
            result2 = get_cached_structure(path, "TEST-WP")
            assert result1 is result2, (
                "get_cached_structure should return the same object (identity) "
                "for unmodified file"
            )
        finally:
            os.unlink(path)
            invalidate_cache()

    @settings(max_examples=5, deadline=None)
    @given(field_id=_field_id_st)
    def test_modified_file_returns_new_object(self, field_id: str):
        """After file modification (mtime changes), get_cached_structure re-parses."""
        invalidate_cache()

        doc = Document()
        doc.add_paragraph(f"${{{field_id}:original}}")

        fd, path = tempfile.mkstemp(suffix=".docx")
        os.close(fd)
        doc.save(path)

        try:
            result1 = get_cached_structure(path, "TEST-WP")

            # Wait to ensure mtime difference on Windows (1-second resolution)
            time.sleep(1.1)
            doc2 = Document()
            doc2.add_paragraph(f"${{{field_id}:modified}}")
            doc2.save(path)

            result2 = get_cached_structure(path, "TEST-WP")
            assert result1 is not result2, (
                "get_cached_structure should return a NEW object after file modification"
            )
        finally:
            os.unlink(path)
            invalidate_cache()


# ─── Property 4: Placeholder extraction round-trip ───────────────────────────


class TestProperty4RoundTrip:
    """Property 4: Placeholder extraction round-trip.

    format_placeholder_summary → re-parse same file → same field_ids and count.

    **Validates: Requirements 2.6**
    """

    @settings(max_examples=5)
    @given(placeholders=placeholder_list_strategy())
    def test_roundtrip_same_field_ids(self, placeholders: list[tuple[str, str]]):
        """Parse→format_summary→re-parse same file produces same field_ids."""
        path = _create_docx_with_placeholders(placeholders)
        try:
            result1 = parse_template(path)
            summary = format_placeholder_summary(result1)
            assert isinstance(summary, str)

            # Re-parse same file
            result2 = parse_template(path)

            ids1 = [p.field_id for p in result1.placeholders]
            ids2 = [p.field_id for p in result2.placeholders]
            assert ids1 == ids2, (
                f"Round-trip field_ids mismatch: first={ids1}, second={ids2}"
            )
        finally:
            os.unlink(path)

    @settings(max_examples=5)
    @given(placeholders=placeholder_list_strategy())
    def test_roundtrip_same_count(self, placeholders: list[tuple[str, str]]):
        """Parse→format_summary→re-parse same file produces same placeholder count."""
        path = _create_docx_with_placeholders(placeholders)
        try:
            result1 = parse_template(path)
            _ = format_placeholder_summary(result1)
            result2 = parse_template(path)
            assert len(result1.placeholders) == len(result2.placeholders), (
                f"Placeholder count mismatch: first={len(result1.placeholders)}, "
                f"second={len(result2.placeholders)}"
            )
        finally:
            os.unlink(path)

    @settings(max_examples=5)
    @given(placeholders=placeholder_list_strategy())
    def test_summary_format_contains_all_fields(self, placeholders: list[tuple[str, str]]):
        """format_placeholder_summary output contains all field_ids from the structure."""
        path = _create_docx_with_placeholders(placeholders)
        try:
            result = parse_template(path)
            summary = format_placeholder_summary(result)
            for p in result.placeholders:
                assert p.field_id in summary, (
                    f"field_id '{p.field_id}' not found in summary output"
                )
        finally:
            os.unlink(path)

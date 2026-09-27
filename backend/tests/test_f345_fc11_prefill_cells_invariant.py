# -*- coding: utf-8 -*-
"""FC-11 判据：prefill 预设的公式格只用 `cells`，工具链不再造 `items`。

spec:
  f3-sync-coverage-and-first-canary · Task 5 / Task 17 · Requirement 7.2 · Property 12
  f4-sync-coverage-and-first-canary · Task 4 / Task 18 · Requirement 7.2 · Property 15
  f5-sync-coverage-and-first-canary · Task 4 / Task 19 · Requirement 7.2 · Property 21

FC-11 原文（F1 design §FC-11）：全库 mapping 块中有 7 个块 `cells` 为空、公式写在 `items` 里，
全部属 F3/F4/F5。运行时四个消费方**全部只读 `cells`**（按值实测）：

  * `wp_template_init_service.py:670`             预填写入 xlsx
  * `formula_management/preset_library.py:163`    公式管理页预设
  * `formula_reverse_index.py:294`                反向索引建边
  * `linkage_graph_builder.py:214`                依赖图

⇒ 这些公式在公式管理页看不到、不预填、不建边。

🔴 根因在工具链（实测比 spec 记载多一处）：`fix_f_cycle_prefill_presets.py` 的
`_ensure_cells` 与 `_ensure_block` 两处都会主动创建 `items` 键，且 `--check` 兼容读 `items`
⇒ 脚本报「0 项欠账」而运行时一条都不生效。两处已修，本文件守护不回退。

🔴 F5 是唯一能用「从 0 到有」证明修复生效的科目：`workpaper:F5` 预设数 0 → 14。
   F1/F2/F3/F4 的预设本来就非零，只能证明「不变差」。
"""
from __future__ import annotations

import collections
import glob
import json
import os
import subprocess
import sys
from pathlib import Path

import pytest

_BACKEND = Path(__file__).resolve().parent.parent
MAPPING_PATH = _BACKEND / "data" / "prefill_formula_mapping.json"
FIX_SCRIPT = _BACKEND / "scripts" / "fix" / "fix_f_cycle_prefill_presets.py"
TEMPLATE_DIR = _BACKEND / "wp_templates"

#: 迁移后各科目的预设条数（实测锚点）。红基线：F3=18 / F4=18 / F5 缺席。
EXPECTED_MIN_PRESETS = {"F3": 22, "F4": 22, "F5": 14}

#: F5 唯一的预设块（FC-11 的 14 条公式全在这里）。
F5_BLOCK_SHEET = "营业务成本审定表F5-1"


@pytest.fixture(scope="module")
def doc() -> dict:
    return json.loads(MAPPING_PATH.read_text(encoding="utf-8"))


@pytest.fixture(scope="module")
def preset_counts() -> collections.Counter:
    """`convert_prefill_presets()` 现算的 page_key → 条数（运行时真实口径）。"""
    sys.path.insert(0, str(_BACKEND))
    from app.services.formula_management.preset_library import convert_prefill_presets

    return collections.Counter(e.page_key for e in convert_prefill_presets())


@pytest.fixture(scope="module")
def visible_tabs() -> dict:
    """F3/F4/F5 源 xlsx 的可见 sheet tab 名（openpyxl 直读）。"""
    try:
        from openpyxl import load_workbook
    except ImportError:  # pragma: no cover
        pytest.skip("openpyxl not available")

    tabs: dict[str, set[str]] = {}
    for code in ("F3", "F4", "F5"):
        pattern = os.path.join(TEMPLATE_DIR, "**", f"{code} *.xlsx")
        files = [f for f in glob.glob(pattern, recursive=True) if "~$" not in f]
        if not files:
            continue
        wb = load_workbook(files[0], read_only=True, data_only=True)
        try:
            tabs[code] = {sn for sn in wb.sheetnames if wb[sn].sheet_state == "visible"}
        finally:
            wb.close()
    return tabs


def _blocks_of(doc: dict, wp_code: str) -> list[dict]:
    return [b for b in doc["mappings"] if b.get("wp_code") == wp_code]


# ═══════════════════════════════════════════════════════════════════════════
# 子判据 ①：全库无 `items` 键（F3-P12① / F4-P15① / F5-P21①）
# ═══════════════════════════════════════════════════════════════════════════


class TestNoItemsKey:
    @pytest.mark.parametrize("wp_code", ["F3", "F4", "F5"])
    def test_no_items_key_per_code(self, doc, wp_code):
        offenders = [
            f"{wp_code}/{b.get('sheet')!r}: items={len(b.get('items') or [])} 条"
            for b in _blocks_of(doc, wp_code)
            if "items" in b
        ]
        assert not offenders, (
            f"{wp_code} 仍有 {len(offenders)} 个块带 `items` 键（运行时死配置，FC-11）:\n"
            + "\n".join(offenders)
        )

    def test_no_items_key_whole_file(self, doc):
        """🔴 全库范围 —— FC-11 的 7 个 items 块全属 F3/F4/F5，迁完应为 0。"""
        offenders = [
            f"[{i}] {b.get('wp_code')}/{b.get('sheet')!r}"
            for i, b in enumerate(doc["mappings"])
            if "items" in b
        ]
        assert not offenders, (
            f"全库仍有 {len(offenders)} 个块带 `items` 键:\n" + "\n".join(offenders)
        )


# ═══════════════════════════════════════════════════════════════════════════
# 子判据 ②：工具链根因两处都不再造 `items`（F3-P12②）
# ═══════════════════════════════════════════════════════════════════════════


class TestToolchainRootCause:
    """🔴 spec 的 Requirement 7.2③ 只点了 `_ensure_cells`；实测 `_ensure_block` 的
    `new_block` 同样造 `items`。两处都要守，否则新增块仍产出死配置。
    """

    @staticmethod
    def _items_write_sites(source: str) -> list[str]:
        """走 AST 找**真实代码**里创建/写入 `items` 键的位置。

        🔴 不用文本匹配：本脚本的 docstring 与注释里逐字引用了修复前的旧写法
        （`block["items"] = []` / `items_key = "items" if ...`）作为根因说明，
        文本匹配会把说明文字当成残留代码，判据自己就假红。AST 只看语法结构。
        """
        import ast

        tree = ast.parse(source)
        sites: list[str] = []
        for node in ast.walk(tree):
            # ① X["items"] = ...
            if isinstance(node, ast.Assign):
                for target in node.targets:
                    if (
                        isinstance(target, ast.Subscript)
                        and isinstance(target.slice, ast.Constant)
                        and target.slice.value == "items"
                    ):
                        sites.append(f'L{node.lineno}: 下标赋值 ...["items"] = ...')
            # ② dict 字面量含 "items" 键
            if isinstance(node, ast.Dict):
                for key in node.keys:
                    if isinstance(key, ast.Constant) and key.value == "items":
                        sites.append(f"L{node.lineno}: dict 字面量含 'items' 键")
            # ③ X.setdefault("items", ...)
            if isinstance(node, ast.Call) and isinstance(node.func, ast.Attribute):
                if node.func.attr == "setdefault" and node.args:
                    first = node.args[0]
                    if isinstance(first, ast.Constant) and first.value == "items":
                        sites.append(f"L{node.lineno}: setdefault('items', ...)")
        return sites

    @staticmethod
    def _items_key_vars(source: str) -> list[str]:
        """走 AST 找 `items_key = ...` 这类兼容读变量赋值。"""
        import ast

        tree = ast.parse(source)
        found: list[str] = []
        for node in ast.walk(tree):
            if isinstance(node, ast.Assign):
                for target in node.targets:
                    if isinstance(target, ast.Name) and target.id.endswith("items_key"):
                        found.append(f"L{node.lineno}: {target.id} = ...")
        return found

    def test_source_has_no_items_creation(self):
        source = FIX_SCRIPT.read_text(encoding="utf-8")
        sites = self._items_write_sites(source)
        assert not sites, (
            "工具链仍会创建 `items` 键（FC-11 根因未修净）:\n"
            + "\n".join(f"  {s}" for s in sites)
        )

    def test_no_items_key_compat_read(self):
        """`items_key = "items" if ... else "cells"` 兼容读必须清零。"""
        source = FIX_SCRIPT.read_text(encoding="utf-8")
        found = self._items_key_vars(source)
        assert not found, (
            "仍有 items_key 兼容读变量（会让死配置重新显绿）:\n"
            + "\n".join(f"  {s}" for s in found)
        )

    def test_single_cells_accessor_exists(self):
        """单一口径 helper `_cells()` 必须存在（否则兼容读会重新长回来）。"""
        import ast

        source = FIX_SCRIPT.read_text(encoding="utf-8")
        tree = ast.parse(source)
        names = {
            node.name for node in ast.walk(tree) if isinstance(node, ast.FunctionDef)
        }
        assert "_cells" in names, "缺少 `_cells()` 单一口径 helper"

    def test_mutation_detects_reintroduced_items_creation(self):
        """反向自检：判据必须能抓住重新引入的 `items` 创建（否则是空转守卫）。"""
        mutated = 'def _ensure_block(doc, d):\n    return {"wp_code": d["wp_code"], "items": []}\n'
        assert self._items_write_sites(mutated), "AST 判据抓不到 dict 字面量里的 items 键"
        mutated2 = 'def f(b):\n    b["items"] = []\n'
        assert self._items_write_sites(mutated2), "AST 判据抓不到下标赋值形态"
        mutated3 = 'def f(b):\n    items_key = "items" if "items" in b else "cells"\n'
        assert self._items_key_vars(mutated3), "AST 判据抓不到 items_key 兼容读"


# ═══════════════════════════════════════════════════════════════════════════
# 子判据 ③：`--check` 对「块只有 items」必红（变异判据，F3-P12③）
# ═══════════════════════════════════════════════════════════════════════════


class TestCheckRejectsItemsBlock:
    """🔴 改造前 `--check` 兼容读 `items` ⇒ 死配置也报 exit 0。本判据把它钉死。"""

    def test_check_passes_on_current_data(self):
        result = subprocess.run(
            [sys.executable, str(FIX_SCRIPT), "--check"],
            capture_output=True,
            text=True,
            encoding="utf-8",
            errors="replace",
            cwd=str(_BACKEND.parent),
        )
        assert result.returncode == 0, (
            f"--check 在当前数据上应 exit 0:\n{result.stdout}\n{result.stderr}"
        )

    def test_check_fails_when_block_reverts_to_items(self, tmp_path, monkeypatch):
        """变异：把 F5 块改回 items 型 ⇒ `validate()` 必红。"""
        sys.path.insert(0, str(_BACKEND / "scripts" / "fix"))
        import importlib

        module = importlib.import_module("fix_f_cycle_prefill_presets")

        mutated = json.loads(MAPPING_PATH.read_text(encoding="utf-8"))
        target = next(
            b for b in mutated["mappings"]
            if b.get("wp_code") == "F5" and b.get("sheet") == F5_BLOCK_SHEET
        )
        target["items"] = target.pop("cells")

        issues = module.validate(mutated)
        assert issues, "把 F5 块改回 items 型后 validate() 仍报 0 项欠账 —— 判据空转"
        joined = "\n".join(issues)
        assert "items" in joined and "F5" in joined, (
            f"欠账消息未指出 F5 的 items 残留:\n{joined}"
        )

    def test_check_fails_on_unknown_sheet_name(self):
        """变异：把 F3 块的 sheet 名改成模板里不存在的 ⇒ 必红。"""
        sys.path.insert(0, str(_BACKEND / "scripts" / "fix"))
        import importlib

        module = importlib.import_module("fix_f_cycle_prefill_presets")

        mutated = json.loads(MAPPING_PATH.read_text(encoding="utf-8"))
        target = next(
            b for b in mutated["mappings"]
            if b.get("wp_code") == "F3" and b.get("sheet") == "附注披露信息(国企)"
        )
        target["sheet"] = "附注披露信息（国企）"  # 全角 —— 模板真名是半角

        issues = module.validate(mutated)
        assert any("可见 tab" in i for i in issues), (
            f"全角 sheet 名未被 `--check` 拦住:\n" + "\n".join(issues)
        )


# ═══════════════════════════════════════════════════════════════════════════
# 子判据 ④：运行时预设条数（F3-P12④ / F5-P21②）
# ═══════════════════════════════════════════════════════════════════════════


class TestRuntimePresetCounts:
    """`convert_prefill_presets()` 是公式管理页的真实数据源 —— 只读 `cells`。

    红基线（迁移前实测）：`workpaper:F0=2 / F1=33 / F2=78 / F3=18 / F4=18`，**F5 缺席**。
    """

    @pytest.mark.parametrize("wp_code", ["F3", "F4", "F5"])
    def test_preset_count_reaches_expected(self, preset_counts, wp_code):
        page_key = f"workpaper:{wp_code}"
        actual = preset_counts.get(page_key, 0)
        expected = EXPECTED_MIN_PRESETS[wp_code]
        assert actual >= expected, (
            f"{page_key} 预设 {actual} 条，期望 ≥{expected}（FC-11 迁移后）"
        )

    def test_f5_presets_exactly_fourteen(self, preset_counts):
        """🔴 F5-P21 核心判据：F5 从 **0** 变 **14**（唯一能证明 FC-11 生效的科目）。"""
        actual = preset_counts.get("workpaper:F5", 0)
        assert actual == 14, (
            f"workpaper:F5 预设 {actual} 条，期望恰好 14（块 `{F5_BLOCK_SHEET}` 的 14 条公式）。"
            "红基线是 0 —— 该数字是 FC-11 修复生效的唯一正面证明"
        )

    def test_non_f345_codes_unchanged(self, preset_counts):
        """零回归：F0/F1/F2 的预设数不受本次迁移影响。"""
        assert preset_counts.get("workpaper:F0", 0) == 2
        assert preset_counts.get("workpaper:F1", 0) == 33
        assert preset_counts.get("workpaper:F2", 0) == 78


# ═══════════════════════════════════════════════════════════════════════════
# 子判据 ⑤：F5 块形态与 PREV sheet 名（F5-P21①③）
# ═══════════════════════════════════════════════════════════════════════════


class TestF5BlockShape:
    def test_f5_block_has_fourteen_cells(self, doc):
        blocks = _blocks_of(doc, "F5")
        assert len(blocks) == 1, f"F5 应恰有 1 个预设块，实得 {len(blocks)}"
        block = blocks[0]
        assert block.get("sheet") == F5_BLOCK_SHEET
        assert "items" not in block
        assert len(block.get("cells") or []) == 14

    def test_f5_prev_references_real_sheet_name(self, doc, visible_tabs):
        """🔴 块内 `PREV('F5', <sheet>, …)` 的 sheet 名必须是模板真名。

        模板真名带「营业务」错字（`营业务成本审定表F5-1`）—— 改 sheet 名会动 sha256 与
        135 处跨表引用，属模板治理债 ⇒ 本判据反向要求 prefill 配置**对齐错字**。
        """
        if "F5" not in visible_tabs:
            pytest.skip("F5 源 xlsx 未找到")
        tabs = visible_tabs["F5"]
        block = _blocks_of(doc, "F5")[0]
        bad = []
        for cell in block.get("cells") or []:
            formula = cell.get("formula") or ""
            if "PREV(" not in formula:
                continue
            # 形如 =PREV('F5','<sheet>','<cell_ref>')
            parts = formula.split("'")
            if len(parts) < 4:
                continue
            sheet = parts[3]
            if sheet not in tabs:
                bad.append(f"{cell.get('cell_ref')}: PREV 引用 {sheet!r} 不在模板 tab")
        assert not bad, "\n".join(bad) + f"\n可用 tab: {sorted(tabs)}"


# ═══════════════════════════════════════════════════════════════════════════
# 子判据 ⑥：全角重复块已删（F4-P15② / F3 需求 7.2②）
# ═══════════════════════════════════════════════════════════════════════════


class TestFullwidthDuplicateBlocksRemoved:
    @pytest.mark.parametrize("wp_code", ["F3", "F4"])
    def test_fullwidth_disclosure_block_gone(self, doc, wp_code):
        fullwidth = [
            b for b in _blocks_of(doc, wp_code)
            if b.get("sheet") == "附注披露信息（国企）"
        ]
        assert not fullwidth, (
            f"{wp_code} 的全角重复块 `附注披露信息（国企）` 未删除 —— "
            "模板真名是半角，全角块运行时解析不到 sheet"
        )

    @pytest.mark.parametrize("wp_code", ["F3", "F4"])
    def test_halfwidth_disclosure_block_kept(self, doc, wp_code):
        halfwidth = [
            b for b in _blocks_of(doc, wp_code)
            if b.get("sheet") == "附注披露信息(国企)"
        ]
        assert len(halfwidth) == 1, (
            f"{wp_code} 的半角块 `附注披露信息(国企)` 应恰有 1 个，实得 {len(halfwidth)}"
        )
        assert len(halfwidth[0].get("cells") or []) >= 2

    @pytest.mark.parametrize("wp_code", ["F3", "F4", "F5"])
    def test_all_sheet_names_in_template_tabs(self, doc, visible_tabs, wp_code):
        """所有块的 sheet 名都必须是源 xlsx 可见 tab（F4-P15③ 的等价断言）。"""
        if wp_code not in visible_tabs:
            pytest.skip(f"{wp_code} 源 xlsx 未找到")
        tabs = visible_tabs[wp_code]
        bad = [
            f"{wp_code}/{b.get('sheet')!r}"
            for b in _blocks_of(doc, wp_code)
            if b.get("sheet") and b["sheet"] not in tabs
        ]
        assert not bad, "\n".join(bad) + f"\n可用 tab: {sorted(tabs)}"

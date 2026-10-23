"""准则转换生成「目标准则独有底稿」用真实模板（spec formula-fetch-and-wp-generation-defect-closure R5.3）。

与按编码生成同款缺陷：``WpConversionGenerateMixin._copy_template_file`` 按 ``gt_template_library.json`` 精确匹配，
E1 这类多页底稿不在库里 ⇒ 落成三行空白册。现与链路编排同一解析（``wp_generate_file.materialize_workpaper_file``）。
SQLite 真 ORM + 真实模板目录，底稿存储目录指到临时目录。
"""
from __future__ import annotations

import uuid

import openpyxl
import pytest
import sqlalchemy as sa
from sqlalchemy.dialects.sqlite.base import SQLiteTypeCompiler
from sqlalchemy.ext.asyncio import AsyncSession, async_sessionmaker, create_async_engine
from sqlalchemy.pool import StaticPool

SQLiteTypeCompiler.visit_JSONB = SQLiteTypeCompiler.visit_JSON

import app.models  # noqa: E402, F401
from app.models.base import Base, ProjectStatus  # noqa: E402
from app.models.core import Project, User  # noqa: E402
from app.models.workpaper_models import WorkingPaper, WpIndex  # noqa: E402
from app.services import wp_generate_file, wp_template_init_service  # noqa: E402
from app.services.wp_generate_file import BLANK_NOTICE, BLANK_NOTICE_ROW  # noqa: E402
from app.services.wp_standard_conversion_service import WpStandardConversionService  # noqa: E402

pytestmark = pytest.mark.asyncio

_TABLES = [User.__table__, Project.__table__, WpIndex.__table__, WorkingPaper.__table__]


@pytest.fixture
async def world(tmp_path, monkeypatch):
    monkeypatch.chdir(tmp_path)
    storage = tmp_path / "wp_storage"
    monkeypatch.setattr(wp_template_init_service, "STORAGE_DIR", storage)
    engine = create_async_engine("sqlite+aiosqlite:///:memory:", poolclass=StaticPool)
    async with engine.begin() as conn:
        await conn.run_sync(lambda c: Base.metadata.create_all(c, tables=_TABLES))
    async with async_sessionmaker(engine, class_=AsyncSession, expire_on_commit=False)() as db:
        pid = uuid.uuid4()
        db.add(Project(id=pid, name="【测试】准则转换生成", client_name="【测试】准则转换生成",
                       status=ProjectStatus.execution, template_type="listed", report_scope="standalone",
                       audit_year=2025, scenario="normal"))
        await db.commit()
        yield {"db": db, "pid": pid, "storage": storage, "root": tmp_path}
    await engine.dispose()


async def _papers(db, pid):
    rows = (await db.execute(
        sa.select(WpIndex.wp_code, WpIndex.wp_name, WorkingPaper.file_path, WorkingPaper.id)
        .join(WorkingPaper, WorkingPaper.wp_index_id == WpIndex.id)
        .where(WpIndex.project_id == pid).order_by(WpIndex.wp_code)
    )).all()
    return {r[0]: r for r in rows}


async def _convert(world, codes):
    svc = WpStandardConversionService(world["db"])
    return await svc._create_target_only_workpapers(
        world["pid"], codes, {"entity_type": "listed", "scope": "standalone", "year": 2025}, None)


async def test_target_only_codes_get_real_template_and_source_name(world):
    assert await _convert(world, ["E1", "ZZ9"]) == 2
    papers = await _papers(world["db"], world["pid"])
    code, name, file_path, wp_id = papers["E1"]
    assert name == "货币资金"
    assert file_path == str(world["storage"] / str(world["pid"]) / f"{wp_id}.xlsx")
    wb = openpyxl.load_workbook(file_path, read_only=True)
    try:
        sheets = wb.sheetnames
    finally:
        wb.close()
    assert "货币资金审定表E1-1" in sheets and "银行存款截止测试表E1-21" in sheets
    assert "现金交易分析E1-26" not in sheets, "normal 场景不合并 IPO 文件"
    blank = openpyxl.load_workbook(papers["ZZ9"][2])
    try:
        assert blank.active[f"A{BLANK_NOTICE_ROW}"].value == BLANK_NOTICE
    finally:
        blank.close()
    assert not (world["root"] / "storage").exists(), "不应再往旧的 storage/projects 相对路径写空白册"


async def test_failure_of_one_code_rolls_back_only_that_code(world, monkeypatch):
    original = wp_generate_file.materialize_workpaper_file

    def _boom(project_id, wp_id, wp_code, wp_name, **kw):
        if wp_code == "H1":
            raise RuntimeError("模板复制失败")
        return original(project_id, wp_id, wp_code, wp_name, **kw)

    monkeypatch.setattr(wp_generate_file, "materialize_workpaper_file", _boom)
    assert await _convert(world, ["H1", "E1"]) == 1
    papers = await _papers(world["db"], world["pid"])
    assert sorted(papers) == ["E1"], "失败的 H1 不应留下无文件的底稿行"
    idx = (await world["db"].execute(sa.select(WpIndex.wp_code).where(WpIndex.project_id == world["pid"]))).scalars()
    assert sorted(idx) == ["E1"]

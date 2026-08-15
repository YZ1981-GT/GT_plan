"""G7 源 xlsx 列事实投影守卫（Task 2：后端事实投影 stale 守卫）。

被测对象（只读，本守卫不改动它们）
----------------------------------
* ``backend/scripts/diagnose/diagnose_g7_column_alignment.py`` —— 只读诊断脚本
* ``backend/data/g7_column_source_facts.json`` —— 38 张表的源 xlsx 列事实投影

守卫的四件事
------------
1. **stale 检测（核心）** —— ``build_facts_payload(read_source_facts())`` 与落盘 JSON
   逐字相等（忽略 ``_meta.generated_at`` 这类时间戳，但 ``_meta.source_sha256`` 必须相等）。
   源 xlsx 一改、投影没重跑，立刻打红并给出重跑命令。
2. **扫描面反向自检** —— 防「正则失效 / 表区注册表被删空」后守卫空转：variant 数、
   表总数、每表列数、label_header 非空、``REGIONS`` 条目数与落盘表数一致。
3. **只读性源码断言** —— 剥注释（``tokenize`` 剥 ``#`` + ``ast`` 剥 docstring，
   **不剥普通字符串字面量**）后脚本代码里不得出现破坏性写操作；配「剥注释确实生效」自检。
4. **篡改必打红** —— 在 ``tmp_path`` 的副本上改一个格，断言比对函数判不相等
   （绝不改动真实落盘文件）。

判据单一真源
------------
「逐字相等」这条不变式只由本文件的 ``diff_payloads()`` 实现一次，正向比对与篡改检验
共用它 —— 否则改一处另一处不红。

Usage（一律从仓库根跑）::

    python -m pytest backend/tests/four_table/test_g7_column_source_facts.py -q

spec: .kiro/specs/g7-column-alignment-and-extraction-closure/ (Task 2)
"""
from __future__ import annotations

import ast
import copy
import hashlib
import importlib.util
import io
import json
import re
import sys
import tokenize
from functools import lru_cache
from pathlib import Path
from typing import Any

import pytest

# backend/tests/four_table/x.py -> _HERE.parents[1] == backend/ ; parents[2] == 仓库根
_HERE = Path(__file__).resolve().parent
_BACKEND = _HERE.parents[1]
_REPO_ROOT = _HERE.parents[2]

SCRIPT_PATH = _BACKEND / "scripts" / "diagnose" / "diagnose_g7_column_alignment.py"
FACTS_PATH_EXPECTED = _BACKEND / "data" / "g7_column_source_facts.json"

EMIT_CMD = (
    "python backend/scripts/diagnose/diagnose_g7_column_alignment.py --emit-facts"
)
RERUN_HINT = f"请重跑 `{EMIT_CMD}`"

# 扫描面下限（只许上调，不许下调 —— 下调等于让守卫在真源缩水时空转）
_MIN_TABLES_TOTAL = 38
_MIN_TABLES_LISTED = 15
_MIN_TABLES_SOE = 23
_MIN_COLUMNS_PER_TABLE = 2
_EXPECTED_VARIANTS = {"listed", "soe"}

# 忽略比对的 _meta 键（时间戳类）。source_sha256 **不得**进这里 —— 它是源模板被换掉的
# 唯一判据；有一条断言专门钉住这一点。
_VOLATILE_META_KEYS: tuple[str, ...] = ("generated_at",)


# ─────────────────────────── 被测脚本动态加载 ───────────────────────────
#
# 该脚本在 backend/scripts/diagnose/ 下不是包，只能 spec_from_file_location 动态加载。
# 🔴 必须先 sys.modules[name] = mod 再 exec_module：否则模块内的 @dataclass 在
#    dataclasses._is_type 里取 sys.modules.get(cls.__module__).__dict__ 得 None，
#    抛 AttributeError（与被测代码无关，极易误判成脚本有语法问题）。
_MODULE_NAME = "_g7_diagnose_under_test"


@lru_cache(maxsize=1)
def diag_module() -> Any:
    assert SCRIPT_PATH.exists(), f"被测脚本缺失：{SCRIPT_PATH}"
    spec = importlib.util.spec_from_file_location(_MODULE_NAME, SCRIPT_PATH)
    assert spec is not None and spec.loader is not None, f"无法为 {SCRIPT_PATH} 构造 spec"
    mod = importlib.util.module_from_spec(spec)
    sys.modules[_MODULE_NAME] = mod  # 必须在 exec_module 之前
    spec.loader.exec_module(mod)
    return mod


@lru_cache(maxsize=1)
def live_payload() -> dict[str, Any]:
    """实时 openpyxl 读源 xlsx 并构造投影载荷。"""
    mod = diag_module()
    return mod.build_facts_payload(mod.read_source_facts())


@lru_cache(maxsize=1)
def disk_payload() -> dict[str, Any]:
    assert FACTS_PATH_EXPECTED.exists(), (
        f"落盘投影缺失：{FACTS_PATH_EXPECTED}。{RERUN_HINT}"
    )
    return json.loads(FACTS_PATH_EXPECTED.read_text(encoding="utf-8"))


# ─────────────────────────── 判据单一真源：逐字相等 ───────────────────────────


def _strip_volatile(payload: Any) -> Any:
    out = copy.deepcopy(payload)
    meta = out.get("_meta") if isinstance(out, dict) else None
    if isinstance(meta, dict):
        for key in _VOLATILE_META_KEYS:
            meta.pop(key, None)
    return out


def _walk(path: str, live: Any, disk: Any) -> list[str]:
    if isinstance(live, dict) and isinstance(disk, dict):
        diffs: list[str] = []
        for key in sorted(set(live) | set(disk), key=str):
            if key not in live:
                diffs.append(f"{path}.{key}: 落盘多出该键（实时读取里没有）")
            elif key not in disk:
                diffs.append(f"{path}.{key}: 落盘缺少该键（实时读取里有）")
            else:
                diffs.extend(_walk(f"{path}.{key}", live[key], disk[key]))
        return diffs
    if isinstance(live, list) and isinstance(disk, list):
        if len(live) != len(disk):
            return [f"{path}: 长度不同 实时={len(live)} 落盘={len(disk)}"]
        diffs = []
        for idx, (a, b) in enumerate(zip(live, disk)):
            diffs.extend(_walk(f"{path}[{idx}]", a, b))
        return diffs
    if type(live) is not type(disk):
        return [
            f"{path}: 类型不同 实时={type(live).__name__} 落盘={type(disk).__name__}"
        ]
    if live != disk:
        return [f"{path}: 值不同 实时={live!r} 落盘={disk!r}"]
    return []


def diff_payloads(live: Any, disk: Any) -> list[str]:
    """比对两份投影载荷，返回差异描述列表（空列表 = 逐字相等）。

    正向 stale 检测与篡改检验共用本函数，避免同一不变式两处各写一份。
    """
    return _walk("$", _strip_volatile(live), _strip_volatile(disk))


def _stale_message(diffs: list[str]) -> str:
    head = diffs[:12]
    more = "" if len(diffs) <= 12 else f"\n... 另有 {len(diffs) - 12} 处差异未列出"
    return (
        "落盘的 g7_column_source_facts.json 与实时 openpyxl 读取不一致（stale）。\n"
        f"{RERUN_HINT}\n"
        "若源模板 G7 长期股权投资.xlsx 确有变更，请连同变更依据一并记入 spec。\n"
        "差异明细：\n  " + "\n  ".join(head) + more
    )


# ─────────────────────────── 剥注释（判据真源） ───────────────────────────
#
# 只剥 `#` 注释与 docstring；**普通字符串字面量保留**（字典键名 / SQL 片段 / 路径字面量
# 都可能是真实消费，剥掉会让判据静默失效）。


def _docstring_line_span(node: ast.AST) -> tuple[int, int] | None:
    body = getattr(node, "body", None)
    if not body:
        return None
    first = body[0]
    if (
        isinstance(first, ast.Expr)
        and isinstance(first.value, ast.Constant)
        and isinstance(first.value.value, str)
    ):
        return (first.lineno, first.end_lineno or first.lineno)
    return None


def strip_comments_and_docstrings(src: str) -> str:
    """剥 `#` 注释（tokenize）+ docstring（ast），保留普通字符串字面量。"""
    lines = src.splitlines(keepends=False)
    blank: set[int] = set()
    try:
        tree = ast.parse(src)
    except SyntaxError as exc:  # pragma: no cover - 被测脚本应始终可解析
        raise AssertionError(f"被测脚本无法 ast.parse：{exc}") from exc
    for node in ast.walk(tree):
        if isinstance(
            node, (ast.Module, ast.FunctionDef, ast.AsyncFunctionDef, ast.ClassDef)
        ):
            span = _docstring_line_span(node)
            if span:
                blank.update(range(span[0], span[1] + 1))
    kept = ["" if (i + 1) in blank else line for i, line in enumerate(lines)]
    no_doc = "\n".join(kept)

    out_lines = list(no_doc.splitlines())
    try:
        tokens = list(tokenize.generate_tokens(io.StringIO(no_doc).readline))
    except tokenize.TokenError as exc:  # pragma: no cover
        raise AssertionError(f"被测脚本无法 tokenize：{exc}") from exc
    # 从后往前删，避免列号位移
    for tok in reversed(tokens):
        if tok.type != tokenize.COMMENT:
            continue
        row = tok.start[0] - 1
        col = tok.start[1]
        if 0 <= row < len(out_lines):
            out_lines[row] = out_lines[row][:col]
    return "\n".join(out_lines)


@lru_cache(maxsize=1)
def raw_source() -> str:
    return SCRIPT_PATH.read_text(encoding="utf-8")


@lru_cache(maxsize=1)
def code_source() -> str:
    return strip_comments_and_docstrings(raw_source())


# ═══════════════════════════ 1. stale 检测 ═══════════════════════════


class TestStaleDetection:
    """Property: 落盘投影必须与实时 openpyxl 读取逐字相等。"""

    def test_live_payload_equals_disk_payload(self) -> None:
        diffs = diff_payloads(live_payload(), disk_payload())
        assert diffs == [], _stale_message(diffs)

    def test_source_sha256_participates_in_comparison(self) -> None:
        """反向锁死：source_sha256 不得被塞进忽略清单（那等于关掉源模板换人的判据）。"""
        assert "source_sha256" not in _VOLATILE_META_KEYS, (
            "source_sha256 被加进了 _VOLATILE_META_KEYS。"
            "它是「源模板被换掉」的唯一判据，不得忽略。"
        )
        mutated = copy.deepcopy(disk_payload())
        mutated["_meta"]["source_sha256"] = "0" * 64
        diffs = diff_payloads(live_payload(), mutated)
        assert any("source_sha256" in d for d in diffs), (
            "改掉 source_sha256 后比对函数没有打红，说明该字段实际未参与比对。"
        )

    def test_generated_at_is_ignored(self) -> None:
        """反向自检：只改时间戳必须判相等，否则「忽略时间戳」是空话。"""
        mutated = copy.deepcopy(disk_payload())
        mutated["_meta"]["generated_at"] = "1999-01-01T00:00:00+00:00"
        assert mutated["_meta"]["generated_at"] != disk_payload()["_meta"]["generated_at"], (
            "构造的时间戳与落盘值相同，本条自检空转"
        )
        assert diff_payloads(live_payload(), mutated) == [], (
            "只改 generated_at 却被判不相等，说明时间戳未被正确忽略"
        )

    def test_stale_message_carries_rerun_hint(self) -> None:
        """报错文案必须给出重跑命令（否则读者不知道怎么修）。"""
        msg = _stale_message(["$.listed: 值不同"])
        assert "请重跑" in msg
        assert EMIT_CMD in msg


# ═══════════════════════════ 2. 扫描面反向自检 ═══════════════════════════


def _iter_tables(payload: dict[str, Any]):
    for variant in sorted(k for k in payload if k != "_meta"):
        sections = payload[variant]
        assert isinstance(sections, dict), f"{variant} 不是 dict"
        for section in sorted(sections):
            tables = sections[section]
            assert isinstance(tables, dict), f"{variant}/{section} 不是 dict"
            for table in sorted(tables):
                yield variant, section, table, tables[table]


class TestScanSurface:
    """防「表区注册表失效 / 正则失效」后守卫空转。"""

    def test_variant_keys_are_exactly_listed_and_soe(self) -> None:
        payload = disk_payload()
        variants = {k for k in payload if k != "_meta"}
        assert variants == _EXPECTED_VARIANTS, (
            f"variant 键集不符：期望 {sorted(_EXPECTED_VARIANTS)}，实际 {sorted(variants)}。"
            f"{RERUN_HINT}"
        )

    def test_table_counts_meet_floor(self) -> None:
        payload = disk_payload()
        per_variant = {
            v: sum(len(t) for t in payload[v].values()) for v in _EXPECTED_VARIANTS
        }
        total = sum(per_variant.values())
        assert per_variant["listed"] >= _MIN_TABLES_LISTED, (
            f"listed 表数 {per_variant['listed']} 低于下限 {_MIN_TABLES_LISTED}，"
            f"扫描面缩水。{RERUN_HINT}"
        )
        assert per_variant["soe"] >= _MIN_TABLES_SOE, (
            f"soe 表数 {per_variant['soe']} 低于下限 {_MIN_TABLES_SOE}，"
            f"扫描面缩水。{RERUN_HINT}"
        )
        assert total >= _MIN_TABLES_TOTAL, (
            f"表总数 {total} 低于下限 {_MIN_TABLES_TOTAL}，扫描面缩水。{RERUN_HINT}"
        )

    def test_every_table_has_at_least_two_columns(self) -> None:
        thin = [
            f"{v}/{s}/{t}: 列数 {len(tf.get('columns') or [])}"
            for v, s, t, tf in _iter_tables(disk_payload())
            if len(tf.get("columns") or []) < _MIN_COLUMNS_PER_TABLE
        ]
        assert thin == [], (
            f"以下表的列数低于 {_MIN_COLUMNS_PER_TABLE}，多半是表头行区间或合并区判定失效：\n  "
            + "\n  ".join(thin)
        )

    def test_every_table_has_nonempty_label_header(self) -> None:
        empty = [
            f"{v}/{s}/{t}"
            for v, s, t, tf in _iter_tables(disk_payload())
            if not str(tf.get("label_header") or "").strip()
        ]
        assert empty == [], (
            "以下表的 label_header 为空（标签列表头读取失效）：\n  " + "\n  ".join(empty)
        )

    def test_regions_count_matches_disk_table_count(self) -> None:
        """表区注册表条目数必须与落盘表数一致 —— 删表区时立刻打红。"""
        regions = diag_module().REGIONS
        disk_total = sum(
            len(tables) for v in _EXPECTED_VARIANTS for tables in disk_payload()[v].values()
        )
        assert len(regions) == disk_total, (
            f"REGIONS 条目数 {len(regions)} 与落盘表总数 {disk_total} 不一致。"
            f"表区注册表被增删后必须重跑投影：{RERUN_HINT}"
        )
        assert len(regions) >= _MIN_TABLES_TOTAL, (
            f"REGIONS 只有 {len(regions)} 条，低于下限 {_MIN_TABLES_TOTAL}"
        )

    def test_region_keys_match_disk_keys(self) -> None:
        region_keys = {(r[0], r[1], r[2]) for r in diag_module().REGIONS}
        disk_keys = {(v, s, t) for v, s, t, _tf in _iter_tables(disk_payload())}
        only_region = sorted(region_keys - disk_keys)
        only_disk = sorted(disk_keys - region_keys)
        assert not only_region and not only_disk, (
            f"REGIONS 与落盘键集不一致。\n仅在 REGIONS：{only_region}\n仅在落盘：{only_disk}\n"
            f"{RERUN_HINT}"
        )

    def test_label_key_convention_is_label(self) -> None:
        """平台惯例锁死：标签列 key 用 'label'（G7 硬编码中文 '项目' 是少数派偏离）。"""
        assert diag_module().LABEL_KEY_CONVENTION == "label", (
            "LABEL_KEY_CONVENTION 已偏离平台惯例 'label'（241/266 标签列定义、跨 70 文件）"
        )


# ═══════════════════════════ 3. 只读性源码断言 ═══════════════════════════

# 破坏性写操作字样：出现在**代码**里即打红（docstring / 注释里说明「本脚本无 --apply」是合法的）
_FORBIDDEN_CODE_TOKENS: tuple[str, ...] = (
    "--apply",
    "db.delete(",
    "DELETE FROM",
)

# 唯一允许的写盘接收者（派生投影 JSON）
_ALLOWED_WRITE_RECEIVERS = {"FACTS_PATH"}

_WRITE_CALL_RE = re.compile(r"([A-Za-z_][\w\.]*)\s*\.\s*write_(?:text|bytes)\s*\(")


class TestScriptIsReadOnly:
    """被测脚本必须保持只读：无 --apply、不删库、不写模板 JSON。"""

    @pytest.mark.parametrize("token", _FORBIDDEN_CODE_TOKENS)
    def test_no_destructive_token_in_code(self, token: str) -> None:
        code = code_source()
        assert token not in code, (
            f"剥注释后脚本代码里出现破坏性字样 {token!r}。"
            "该脚本是只读诊断脚本，不得具备写库 / 应用变更的能力。"
        )

    def test_no_case_insensitive_sql_delete(self) -> None:
        code = code_source()
        hit = re.search(r"\bDELETE\s+FROM\b", code, re.IGNORECASE)
        assert hit is None, f"剥注释后代码里出现 SQL DELETE：{hit.group(0)!r}"

    def test_write_calls_only_target_facts_json(self) -> None:
        code = code_source()
        receivers = set(_WRITE_CALL_RE.findall(code))
        assert receivers, (
            "代码里找不到任何 write_text/write_bytes 调用, "
            "而 --emit-facts 必须写盘, 说明正则失效, 本条判据已空转"
        )
        illegal = sorted(receivers - _ALLOWED_WRITE_RECEIVERS)
        assert not illegal, (
            f"出现不允许的写盘目标 {illegal}。"
            f"只允许写派生投影 JSON（接收者 {sorted(_ALLOWED_WRITE_RECEIVERS)}）"
        )

    def test_no_write_operation_targets_note_template(self) -> None:
        code = code_source()
        offenders: list[str] = []
        for line_no, line in enumerate(code.splitlines(), start=1):
            if ".write_text" not in line and ".write_bytes" not in line:
                continue
            if "note_template" in line or "TEMPLATE_PATHS" in line:
                offenders.append(f"L{line_no}: {line.strip()}")
        assert offenders == [], (
            "脚本存在对 note_template_*.json 的写操作：\n  " + "\n  ".join(offenders)
        )
        # 反向自检：模板路径确实被脚本读取（否则本条判据的扫描面为空）
        assert "TEMPLATE_PATHS" in code, "代码里没有 TEMPLATE_PATHS，本条判据扫描面为空"
        assert "read_text" in code, "代码里没有 read_text，模板读取路径判定失效"

    def test_comment_stripping_actually_works(self) -> None:
        """剥注释生效自检：原文含 --apply（在 docstring 里），剥后必须没有。"""
        raw = raw_source()
        code = code_source()
        assert "--apply" in raw, (
            "原文里找不到 --apply 字样, 该脚本 docstring 本应写明「不提供 --apply」。"
            "找不到说明扫描面或文件已变，本自检失去意义。"
        )
        assert "--apply" not in code, (
            "剥注释后仍含 --apply，说明剥注释未生效（判据会把说明文字数成真实能力）"
        )
        assert len(code) < len(raw), "剥注释后长度未变小，剥注释未生效"

    def test_stripper_keeps_plain_string_literals(self) -> None:
        """剥注释器 fixture 自检：剥 docstring 与 # 注释，但保留普通字符串字面量。"""
        sample = (
            '"""模块 docstring 里提到 --apply 不算真实能力。"""\n'
            "import os  # 行尾注释里也提到 --apply\n"
            "# 独立注释行 --apply\n"
            "def f():\n"
            '    """函数 docstring --apply"""\n'
            '    return "KEEPME--apply"\n'
        )
        stripped = strip_comments_and_docstrings(sample)
        assert "KEEPME--apply" in stripped, "普通字符串字面量被误剥"
        assert stripped.count("--apply") == 1, (
            f"docstring / 注释未被完全剥掉，剩余 --apply 计数 {stripped.count('--apply')}"
        )
        assert "行尾注释" not in stripped and "独立注释行" not in stripped
        assert "模块 docstring" not in stripped and "函数 docstring" not in stripped


# ═══════════════════════════ 4. 篡改必打红 ═══════════════════════════


def _first_table_path(payload: dict[str, Any]) -> tuple[str, str, str]:
    for variant, section, table, _tf in _iter_tables(payload):
        return variant, section, table
    raise AssertionError("落盘投影里没有任何表，无法做篡改检验")


class TestMutationTurnsRed:
    """在 tmp_path 副本上篡改一格，比对函数必须判不相等（绝不动真实文件）。"""

    def test_mutating_a_column_label_is_detected(self, tmp_path: Path) -> None:
        before_md5 = hashlib.md5(FACTS_PATH_EXPECTED.read_bytes()).hexdigest()

        copy_path = tmp_path / "g7_column_source_facts.json"
        copy_path.write_bytes(FACTS_PATH_EXPECTED.read_bytes())
        mutated = json.loads(copy_path.read_text(encoding="utf-8"))

        variant, section, table = _first_table_path(mutated)
        original = mutated[variant][section][table]["columns"][0]["label"]
        mutated[variant][section][table]["columns"][0]["label"] = original + "X"
        copy_path.write_text(
            json.dumps(mutated, ensure_ascii=False, indent=2) + "\n", encoding="utf-8"
        )

        diffs = diff_payloads(
            live_payload(), json.loads(copy_path.read_text(encoding="utf-8"))
        )
        assert diffs, "篡改列 label 后比对函数仍判相等，stale 判据失效"
        assert any(table in d and "columns" in d for d in diffs), (
            f"差异未指向被篡改的表 {table}，实际差异：{diffs[:5]}"
        )

        after_md5 = hashlib.md5(FACTS_PATH_EXPECTED.read_bytes()).hexdigest()
        assert before_md5 == after_md5, "真实落盘文件被本测试改动了，这是禁止的"

    def test_mutating_label_header_is_detected(self, tmp_path: Path) -> None:
        mutated = copy.deepcopy(disk_payload())
        variant, section, table = _first_table_path(mutated)
        mutated[variant][section][table]["label_header"] += "X"
        diffs = diff_payloads(live_payload(), mutated)
        assert any("label_header" in d for d in diffs), (
            f"篡改 label_header 未被检出，实际差异：{diffs[:5]}"
        )

    def test_dropping_a_table_is_detected(self, tmp_path: Path) -> None:
        mutated = copy.deepcopy(disk_payload())
        variant, section, table = _first_table_path(mutated)
        del mutated[variant][section][table]
        diffs = diff_payloads(live_payload(), mutated)
        assert any(table in d and "缺少" in d for d in diffs), (
            f"删表未被检出，实际差异：{diffs[:5]}"
        )

    def test_adding_a_bogus_table_is_detected(self) -> None:
        mutated = copy.deepcopy(disk_payload())
        variant, section, _table = _first_table_path(mutated)
        mutated[variant][section]["凭空多出的表"] = {"label_header": "x", "columns": []}
        diffs = diff_payloads(live_payload(), mutated)
        assert any("凭空多出的表" in d for d in diffs), (
            f"凭空加表未被检出，实际差异：{diffs[:5]}"
        )


# ═══════════════════════════ 5. JSON 只读性声明 ═══════════════════════════


class TestFactsJsonIsDerivedProjection:
    """防后来者把派生投影当真源手改。"""

    @pytest.mark.parametrize("phrase", ["派生投影，非真源", "禁止手改"])
    def test_meta_note_declares_derived_and_no_manual_edit(self, phrase: str) -> None:
        note = str(disk_payload()["_meta"].get("note") or "")
        assert phrase in note, (
            f"_meta.note 缺少 {phrase!r} 字样。该 JSON 是派生投影，"
            "必须在文件内自述，否则后来者会把它当真源手改。实际 note：" + note
        )

    def test_meta_declares_generator_and_source(self) -> None:
        meta = disk_payload()["_meta"]
        assert "diagnose_g7_column_alignment.py" in str(meta.get("note") or ""), (
            "_meta.note 未写明生成脚本"
        )
        assert "G7" in str(meta.get("generated_from") or ""), (
            f"_meta.generated_from 未指向 G7 源模板：{meta.get('generated_from')!r}"
        )


# ═══════════════════════════ 6. source_sha256 与实际 xlsx 一致 ═══════════════════════════


class TestSourceSha256:
    """源模板被换掉时必须打红。"""

    def test_recomputed_sha256_matches_meta(self) -> None:
        mod = diag_module()
        src: Path = mod.SRC_XLSX
        assert src.exists(), f"源模板缺失：{src}"
        digest = hashlib.sha256()
        with src.open("rb") as fh:
            for chunk in iter(lambda: fh.read(1 << 16), b""):
                digest.update(chunk)
        actual = digest.hexdigest()
        recorded = str(disk_payload()["_meta"].get("source_sha256") or "")
        assert actual == recorded, (
            "源模板 sha256 与落盘投影记录不一致（源 xlsx 已被替换或投影已过期）。\n"
            f"实际={actual}\n落盘={recorded}\n{RERUN_HINT}"
        )
        assert len(recorded) == 64, f"source_sha256 长度异常：{len(recorded)}"

    def test_facts_path_constant_points_to_expected_file(self) -> None:
        assert diag_module().FACTS_PATH.resolve() == FACTS_PATH_EXPECTED.resolve(), (
            "脚本的 FACTS_PATH 与守卫期望的落盘路径不一致，守卫可能在盯错文件"
        )


# ═════════════════════════════════════════════════════════════════════════════
# 源模板**自身缺陷**登记表 + stale 检测（Task 17 / Property 23）
# ═════════════════════════════════════════════════════════════════════════════
#
# ## 与已有三张豁免表的区别（别混淆）
#
# `diagnose_g7_column_alignment.py` 里已有三张登记表，它们全部关于**列结构差异**：
#
#   ① `SOURCE_PLACEHOLDER_SINGLE_SLOT`   源模板只填 1 个实体槽 ⇒ 平台按单槽 flat 实现
#   ② `PLATFORM_CONSTRAINT_EXEMPTIONS`   逐字照抄会渲染崩 ⇒ 必要的等价改写
#   ③ `CORNER_TITLE_EXEMPTIONS`          转角标题，源未给行标识列名
#
# 本表是**第四类、性质不同**：源模板**单元格公式引用错了**或**编号漏了**——
# 不是「结构差异」而是「内容错误」，平台按**修正后的意图**实现。
#
# ## 🔴 为什么必须单独读公式
#
# facts 投影与上面三张表都走 `load_workbook(data_only=True)`（取**值**）。
# 公式引用错误在取值路径上**结构性看不见** —— 值可能恰好都是空/都是 0。
# 故本节自己用 `data_only=False` 读**公式原文**，与既有路径互不干扰。
#
# ## stale 检测语义
#
# 每条登记都断言「源 xlsx 里这个缺陷**仍然存在**」。缺陷一旦被上游修正，
# 断言打红 —— 提醒把该条移出登记表并复核平台实现是否还该保持「修正后的意图」。
# 这是「防豁免变永久盲区」的标准手法（平台 memory 已记）。

_DEFECT_CAP = 5  # 🔴 条目数上限，**只许缩短**


class SourceTemplateDefect:
    """源模板自身缺陷的一条登记（纯数据）。"""

    __slots__ = ("locator", "original", "basis", "intent", "probe")

    def __init__(self, locator: str, original: str, basis: str, intent: str, probe: str) -> None:
        #: 源 xlsx 定位（如 `soe!C260`）
        self.locator = locator
        #: 源模板原文（公式或文字，逐字）
        self.original = original
        #: 判定为缺陷的依据（≥20 字；必须能让后人独立复核）
        self.basis = basis
        #: 本平台按什么意图实现
        self.intent = intent
        #: stale 检测用的探针方法名（本类不执行，仅登记，由测试逐条调用）
        self.probe = probe


#: 🔴 5 条源模板缺陷。每条都由同名 `test_stale_*` 断言「缺陷仍在」。
SOURCE_TEMPLATE_DEFECTS: tuple[SourceTemplateDefect, ...] = (
    SourceTemplateDefect(
        locator="soe!C260/D260",
        original="C260 = ='被投资单位财务信息（合营、联营）G7-5'!C12 ；D260 = …!D12",
        basis=(
            "同一张表（soe 重要联营企业主要财务信息 A257:H269）的表头 C257/E257/G257 分别引用 "
            "G7-5 的 E8/G8/I8 三个主体列，紧邻的 r259「流动资产」也按 E10/F10（第一主体）、"
            "G10/H10（第二主体）、I10/J10（第三主体）正确展开；而 r260「非流动资产」却引用 "
            "**C12/D12** —— C/D 是 G7-5 里**合营企业1**的列，不属本表三个主体中的任何一个。"
            "同表同列口径自相矛盾，判定为源模板错列引用。"
        ),
        intent=(
            "平台不照抄该错误引用。联营 FS 表按「3 主体 ×（期末数/期初数）= 6 个数据列」实现"
            "（`_assoc_columns(ASSOCIATE_FS_SLOT, '期末数', '期初数')`），各主体列由审计师"
            "在 G7-5 录入后经跨表引用带入，不硬绑源模板的单元格地址。"
        ),
        probe="_probe_defect_wrong_column_ref",
    ),
    SourceTemplateDefect(
        locator="soe!D282:D286",
        original=(
            "D282 = ='被投资单位财务信息（合营、联营）G7-5'!I49 ；D283=…!I50 ；"
            "D284=…!I51 ；D285=…!I52 ；D286=…!I53"
        ),
        basis=(
            "r282~r286 属「合营企业：」段（A281），其「本期数」列 C282~C286 引用 G7-5 的 "
            "G42~G46；而「上期数」列 D282~D286 引用 **I49~I53**。I49~I53 恰好是下方"
            "「联营企业：」段（A287）的 D288~D292 所引用的同一批单元格 —— 两段的上期数"
            "**逐格完全相同**。合营与联营的汇总数不可能恒等，判定为源模板把合营段的上期数"
            "错指到了联营行。"
        ),
        intent=(
            "平台按语义实现：不重要合营/联营汇总表的两段各自独立取数"
            "（seed 侧 T6 `不重要合营企业和联营企业的汇总信息` 的 12 行分「合营企业：」"
            "与「联营企业：」两组，各组行值由审计师分别录入），不复制源模板的错误引用。"
        ),
        probe="_probe_defect_joint_uses_associate_rows",
    ),
    SourceTemplateDefect(
        locator="listed!B187:G187",
        original=(
            "B187 = ='被投资单位财务信息（合营、联营）G7-5'!E27（与 B186 **逐字相同**）；"
            "C187~G187 同样与 C186~G186 相同"
        ),
        basis=(
            "r186 是「对联营企业权益投资的账面价值」、r187 是「存在公开报价的权益投资的"
            "公允价值」—— 两个不同科目，源模板却让它们引用**同一批单元格**（E27/F27/G27/"
            "H27/I27/J27）。参照 soe 侧同族两行（r268 → E27、r269 → **E28**）可知 r187 "
            "应为 E28 系列。判定为源模板漏改行号（复制上一行后未改引用）。"
        ),
        intent=(
            "平台把两行作为**独立可录入行**实现（listed 重要联营企业主要财务信息的 17 行行集"
            "里两者各占一行），值由 G7-5 经跨表引用或审计师录入，不共用同一取数地址。"
        ),
        probe="_probe_defect_duplicate_cell_ref",
    ),
    SourceTemplateDefect(
        locator="soe!A209..A314 小节编号序列",
        original="(1) r209 长期股权投资明细 → (3) r225 重要合营企业… → (4) r253 → (5) r279 → (6) r297 → (7) r314",
        basis=(
            "soe `八、18` 章节的小节编号从 (1) 直接跳到 (3)，**缺 (2)**。同章节其余编号"
            "连续（3→4→5→6→7），且 A 列扫描确认全表不存在「（2）」形态的本章节小节标题"
            "⇒ 不是被合并或改写，是漏号。"
        ),
        intent=(
            "平台不按源模板的编号排布表序。`fix_note_g7_soe_structure.EXPECTED` 用**表名**"
            "（T0~T9）而非编号定位，前端 `G7_SOE_DISCLOSURE_SECTIONS` 亦按 `templateTableKey` "
            "索引 ⇒ 缺号不影响任何取数/渲染；登记它是为了让后人核对源模板时不误以为平台漏表。"
        ),
        probe="_probe_defect_missing_subsection_number",
    ),
    SourceTemplateDefect(
        locator="soe!B62:L62",
        original="r62 只有 A62='项  目'，B62~L62 **全空**；而 r63 有 5 组「期末数/本期发生额 + 期初数/上期发生额」（C~L 共 10 列）",
        basis=(
            "「2、主要财务信息：」（A61）表的表头行 r63 铺了 10 个数据列 = 5 个主体 × 2 期，"
            "但主体名所在的 r62 除标签列外**一个名字都没有** ⇒ 源模板留的是空白占位，"
            "等审计师填被投资单位名。这与「表头缺失」不同：列数是明确的，缺的只是名字。"
        ),
        intent=(
            "平台按**动态列**实现（`MINORITY_FS_SLOT` + `buildG7SlotColumns`），"
            "默认名 `公司1..公司5` 仅作初始 seed，审计师可增删改名；列 key 用稳定 "
            "`{slot}_{seq}_{sub}` 不随改名漂移。**不把空白当成 5 个固定空列写死。**"
        ),
        probe="_probe_defect_empty_entity_names",
    ),
)


@lru_cache(maxsize=1)
def _formula_workbook():
    """用 `data_only=False` 读源 xlsx（取**公式原文**）。

    🔴 与 facts 投影的 `data_only=True` 路径**互不干扰** —— 公式引用错误在取值
    路径上结构性看不见，故必须另开一次读取。
    """
    from openpyxl import load_workbook

    mod = diag_module()
    assert mod.SRC_XLSX.exists(), f"源模板缺失：{mod.SRC_XLSX}"
    wb = load_workbook(mod.SRC_XLSX, data_only=False)
    return wb, mod.SHEETS


def _cell_formula(variant: str, ref: str) -> str:
    wb, sheets = _formula_workbook()
    value = wb[sheets[variant]][ref].value
    return "" if value is None else str(value)


# ─── 五个 stale 探针（各返回 `(仍存在?, 实测描述)`）────────────────────────────
#
# 🔴 **每个探针都接受可注入的 `read`**（缺省 = 真读源 xlsx 公式）。
#    存在理由：反向自检必须能**真正调用探针**并验证它「对正确形态返回 False」。
#    改造前反向自检是自己重算一遍判据、不碰探针 ⇒ 把探针体改成 `still = True`
#    仍然全绿（本轮变异检验实测到这个 GREEN）—— 那正是「守卫没覆盖判据本体」。
#    有了注入点，`test_reverse_self_check_probes_can_say_no` 给探针喂「已修好」的
#    公式，探针若写死恒真立刻打红。

_Reader = "Callable[[str, str], str]"


def _probe_defect_wrong_column_ref(read=_cell_formula) -> tuple[bool, str]:
    """soe!C260/D260 是否仍引用 G7-5 的 C/D 列（而非 E/F）。"""
    c260 = read("soe", "C260")
    d260 = read("soe", "D260")
    still = "!C12" in c260 and "!D12" in d260
    return still, f"C260={c260!r} D260={d260!r}"


def _probe_defect_joint_uses_associate_rows(read=_cell_formula) -> tuple[bool, str]:
    """soe 合营段上期数（D282:D286）是否仍与联营段（D288:D292）逐格相同。"""
    joint = [read("soe", f"D{r}") for r in range(282, 287)]
    assoc = [read("soe", f"D{r}") for r in range(288, 293)]
    still = joint == assoc and all(j for j in joint)
    return still, f"合营段={joint} 联营段={assoc}"


def _probe_defect_duplicate_cell_ref(read=_cell_formula) -> tuple[bool, str]:
    """listed r187 是否仍与 r186 引用同一批单元格。"""
    row186 = [read("listed", f"{c}186") for c in "BCDEFG"]
    row187 = [read("listed", f"{c}187") for c in "BCDEFG"]
    still = row186 == row187 and all(v for v in row186)
    return still, f"r186={row186[:2]}… r187={row187[:2]}…"


def _probe_defect_missing_subsection_number(read=_cell_formula) -> tuple[bool, str]:
    """soe `八、18` 段的小节编号是否仍缺 (2)。"""
    nums: list[int] = []
    for row in range(209, 320):
        raw = read("soe", f"A{row}")
        if not raw:
            continue
        m = re.match(r"^[（(](\d+)[）)]", raw.strip())
        if m:
            nums.append(int(m.group(1)))
    still = 1 in nums and 2 not in nums and 3 in nums
    return still, f"A209:A319 的小节编号序列={nums}"


def _probe_defect_empty_entity_names(read=_cell_formula) -> tuple[bool, str]:
    """soe r62 的主体名是否仍全空、而 r63 仍铺了 10 个数据列。"""
    filled_names = [v for v in (read("soe", f"{c}62").strip() for c in "BCDEFGHIJKL") if v]
    filled_headers = [v for v in (read("soe", f"{c}63").strip() for c in "CDEFGHIJKL") if v]
    still = not filled_names and len(filled_headers) == 10
    return still, f"r62 非空主体名={filled_names} / r63 非空表头列数={len(filled_headers)}"


_PROBES = {
    "_probe_defect_wrong_column_ref": _probe_defect_wrong_column_ref,
    "_probe_defect_joint_uses_associate_rows": _probe_defect_joint_uses_associate_rows,
    "_probe_defect_duplicate_cell_ref": _probe_defect_duplicate_cell_ref,
    "_probe_defect_missing_subsection_number": _probe_defect_missing_subsection_number,
    "_probe_defect_empty_entity_names": _probe_defect_empty_entity_names,
}


class TestSourceTemplateDefectRegistry:
    """登记表卫生：条数上限 / 依据齐备 / 探针齐备 / 定位可解析。"""

    def test_registry_is_capped(self):
        """🔴 上限只许缩短 —— 防「对不上就往登记表加一条」变成逃逸阀。"""
        assert len(SOURCE_TEMPLATE_DEFECTS) <= _DEFECT_CAP, (
            f"源模板缺陷登记 {len(SOURCE_TEMPLATE_DEFECTS)} 条 > 上限 {_DEFECT_CAP}"
        )

    def test_every_entry_documents_all_three_fields(self):
        """每条必须写「源模板原文 / 判定依据 / 平台实现意图」三件套。"""
        for d in SOURCE_TEMPLATE_DEFECTS:
            assert d.locator.strip(), "缺 locator"
            assert len(d.original) >= 10, f"{d.locator} 源原文过短：{d.original!r}"
            assert len(d.basis) >= 20, f"{d.locator} 依据过短（{len(d.basis)} 字）"
            assert len(d.intent) >= 20, f"{d.locator} 平台意图过短（{len(d.intent)} 字）"

    def test_locators_name_a_real_variant(self):
        """定位必须指向真实 variant（`soe!` / `listed!`），否则无从复核。"""
        _wb, sheets = _formula_workbook()
        for d in SOURCE_TEMPLATE_DEFECTS:
            head = d.locator.split("!", 1)[0]
            assert head in sheets, f"{d.locator} 的 variant {head!r} 不在 {sorted(sheets)}"

    def test_every_entry_has_a_registered_probe(self):
        """每条都必须有 stale 探针，且探针真实可调（防登记表变纯注释）。"""
        for d in SOURCE_TEMPLATE_DEFECTS:
            assert d.probe in _PROBES, f"{d.locator} 的探针 {d.probe} 未注册"
        # 反面：注册的探针不能多于登记条目（防留下无主探针）
        assert set(_PROBES) == {d.probe for d in SOURCE_TEMPLATE_DEFECTS}, (
            "探针集合与登记条目不同域 ⇒ 有探针无登记或有登记无探针"
        )

    def test_probes_are_distinct(self):
        """五条各用**各自**的探针（防一个探针糊五条）。"""
        probes = [d.probe for d in SOURCE_TEMPLATE_DEFECTS]
        assert len(set(probes)) == len(probes), f"探针重复：{probes}"


class TestSourceTemplateDefectsStillPresent:
    """🔴 stale 检测：缺陷一旦被上游修正即打红，提醒移出登记并复核平台实现。

    这不是「希望缺陷永远存在」，而是**防豁免变永久盲区**：登记的前提是
    「源模板确实有这个错」，前提没了登记就该撤，否则平台会一直按「修正后的意图」
    实现一个已经不存在的偏离，而没人知道。
    """

    def test_all_five_defects_still_reproduce(self):
        stale: list[str] = []
        details: list[str] = []
        for d in SOURCE_TEMPLATE_DEFECTS:
            still, observed = _PROBES[d.probe]()
            details.append(f"{d.locator}: still={still} | {observed}")
            if not still:
                stale.append(f"{d.locator}（探针 {d.probe}）\n      实测：{observed}")
        assert not stale, (
            "以下源模板缺陷在当前源 xlsx 里**已不复现** ⇒ 请从 SOURCE_TEMPLATE_DEFECTS "
            "移出，并复核平台是否还该保持「修正后的意图」：\n   "
            + "\n   ".join(stale)
            + "\n\n全部探针实测：\n   "
            + "\n   ".join(details)
        )

    def test_probes_read_formulas_not_values(self):
        """判据自检：探针读到的必须是**公式原文**（`=…` 开头），不是求值结果。

        少了这一条，若哪天 `_formula_workbook` 被误改成 `data_only=True`，
        全部探针会读到 None/空 → `still` 变 False → 上一条把「读法坏了」误报成
        「缺陷已修好」，方向完全相反。
        """
        c260 = _cell_formula("soe", "C260")
        assert c260.startswith("="), (
            f"soe!C260 读到 {c260!r} 不是公式 ⇒ `_formula_workbook` 可能被改成 "
            "data_only=True，全部 stale 探针会失效"
        )
        b186 = _cell_formula("listed", "B186")
        assert b186.startswith("="), f"listed!B186 读到 {b186!r} 不是公式"

    def test_reverse_self_check_probes_can_say_no(self):
        """🔴 反向自检：**真正调用每个探针**，喂「已修好」的源码形态，必须返回 False。

        改造前这条是自己重算一遍判据、**不碰探针本体** ⇒ 把某个探针写成
        `still = True` 仍然全绿（本轮变异检验实测到这个 GREEN）。现在改为给探针注入
        替身读取器（模拟上游把缺陷修好后的源 xlsx），探针若恒真立刻打红。

        每个替身都只把**该缺陷相关的格**改成正确形态，其余格照实读 —— 这样
        「返回 False」只能来自判据真的生效，而不是替身把一切都变空。
        """

        def patched(overrides: dict[tuple[str, str], str]):
            def read(variant: str, ref: str) -> str:
                key = (variant, ref)
                if key in overrides:
                    return overrides[key]
                return _cell_formula(variant, ref)

            return read

        # ① 缺陷 1 修好 = C260/D260 改引用 E12/F12（与同表 r259 口径一致）
        fixed1 = patched(
            {
                ("soe", "C260"): "='被投资单位财务信息（合营、联营）G7-5'!E12",
                ("soe", "D260"): "='被投资单位财务信息（合营、联营）G7-5'!F12",
            }
        )
        still1, obs1 = _probe_defect_wrong_column_ref(fixed1)
        assert still1 is False, f"缺陷 1 探针对「已修好」仍报 still=True ⇒ 恒真（{obs1}）"

        # ② 缺陷 2 修好 = 合营段上期数改引用 I42~I46（自己那一段），不再与联营段相同
        fixed2 = patched(
            {
                ("soe", f"D{282 + i}"): f"='被投资单位财务信息（合营、联营）G7-5'!I{42 + i}"
                for i in range(5)
            }
        )
        still2, obs2 = _probe_defect_joint_uses_associate_rows(fixed2)
        assert still2 is False, f"缺陷 2 探针对「已修好」仍报 still=True ⇒ 恒真（{obs2}）"

        # ③ 缺陷 3 修好 = listed r187 改引用 E28 系列（参照 soe r269）
        fixed3 = patched(
            {
                ("listed", f"{c}187"): f"='被投资单位财务信息（合营、联营）G7-5'!{c2}28"
                for c, c2 in zip("BCDEFG", "EFGHIJ")
            }
        )
        still3, obs3 = _probe_defect_duplicate_cell_ref(fixed3)
        assert still3 is False, f"缺陷 3 探针对「已修好」仍报 still=True ⇒ 恒真（{obs3}）"

        # ④ 缺陷 4 修好 = 补上一个 (2) 小节标题
        fixed4 = patched({("soe", "A224"): "（2）补上的小节标题"})
        still4, obs4 = _probe_defect_missing_subsection_number(fixed4)
        assert still4 is False, f"缺陷 4 探针对「已修好」仍报 still=True ⇒ 恒真（{obs4}）"

        # ⑤ 缺陷 5 修好 = r62 填上 5 个主体名
        fixed5 = patched({("soe", c + "62"): f"被投资单位{i}" for i, c in enumerate("BDFHJ", 1)})
        still5, obs5 = _probe_defect_empty_entity_names(fixed5)
        assert still5 is False, f"缺陷 5 探针对「已修好」仍报 still=True ⇒ 恒真（{obs5}）"

    def test_reverse_self_check_criteria_have_discriminating_power(self):
        """判据区分力：对**未登记为缺陷**的正确行套同样判据必须不成立。

        与上一条互补：上一条证「探针不恒真」，本条证「判据不是宽到什么都算缺陷」。
        """
        # r259 是同表**正确**展开的行（C259→E10 / D259→F10）
        c259 = _cell_formula("soe", "C259")
        d259 = _cell_formula("soe", "D259")
        assert not ("!C12" in c259 and "!D12" in d259), (
            f"缺陷 1 的判据对正确行 r259 也成立（C259={c259!r}）⇒ 判据无区分力"
        )
        # 缺陷 2 的区分力来自「两段本应不同」：用 C 列（本期数）验证两段确实不同
        joint_c = [_cell_formula("soe", f"C{r}") for r in range(282, 287)]
        assoc_c = [_cell_formula("soe", f"C{r}") for r in range(288, 293)]
        assert joint_c != assoc_c, (
            "合营段与联营段的**本期数**也逐格相同 ⇒ 缺陷 2 的「两段本应不同」前提不成立，"
            "需重新判定该缺陷的性质"
        )

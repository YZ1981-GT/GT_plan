"""守卫：G 循环公式预设科目码正确性、sheet 存在性与披露块覆盖。

锁定 `fix_g_cycle_prefill_presets.py`（审定表块纠偏 + sheet 名改名）+
`fix_g_cycle_disclosure_presets.py`（披露块补全）的结果。

Property 8：审定表块的科目码属于本循环报表行
Property 9：损益类审定表块口径 = 本期发生额
Property 10：每个有披露 sheet 的循环都有对应的披露块
Property 11（2026-09-27 新增）：**每个 G 块的 `sheet` ∈ 对应源 xlsx 真实 tab 名**

🔴 Property 11 是本文件此前缺的那条断言（spec `g-cycle-sync-foundation-and-first-canary`
GC-8 / 红基线 RG-5）。D / E1 / F / I / J / K / L / H0 八个循环都有同款断言，
唯独 G 没有 ⇒ 块 `[169]` `明细分析表G13-2` 与 `[170]` `明细分析表G14-2` 两个错名
（真名 `明细表G13-2` / `明细表G14-2`）长期逃检，7 条预设指向不存在的 sheet、
写 xlsx 的预填消费方找不到 sheet 而静默不填。
口径照既有 `backend/tests/test_i_cycle_formula_presets.py::test_sheet_exists_in_template`
（含它的「历史台账 + stale 检测」形态），**不新造第二套**。
"""
from __future__ import annotations

import json
import re
import subprocess
import sys
from pathlib import Path

import pytest

BACKEND = Path(__file__).resolve().parents[2]
MAPPING_PATH = BACKEND / "data" / "prefill_formula_mapping.json"
DISCLOSURE_SCRIPT = BACKEND / "scripts" / "fix" / "fix_g_cycle_disclosure_presets.py"
TEMPLATE_DIR = BACKEND / "wp_templates" / "G"

# 科目真源（g_cycle_specs.py 实证）
G_ACCOUNT_CODES = {
    "G1": "1101", "G2": "1132", "G3": "1131", "G4": "1504",
    "G5": "1531", "G6": "1506", "G7": "1511", "G8": "1507",
    "G9": "1519", "G10": "2101", "G11": "6111", "G12": "6103",
    "G13": "6101", "G14": "6702",
}

# 损益类循环（口径 = 本期发生额，不是期末余额）
PL_CYCLES = {"G11", "G12", "G13", "G14"}


@pytest.fixture(scope="module")
def g_mappings():
    data = json.loads(MAPPING_PATH.read_text(encoding="utf-8"))
    return [m for m in data.get("mappings", []) if m.get("wp_code", "").startswith("G")]


@pytest.fixture(scope="module")
def template_sheets() -> dict[str, list[str]]:
    """{wp_code: [sheetname, ...]}，跳过 `~$` 锁文件。

    口径逐字照 `test_i_cycle_formula_presets._load_template_sheetnames`。
    """
    import openpyxl

    result: dict[str, list[str]] = {}
    for xlsx_path in sorted(TEMPLATE_DIR.glob("*.xlsx")):
        if xlsx_path.name.startswith("~$"):
            continue
        m = re.match(r"^(G\d+)\s", xlsx_path.stem)
        if not m:
            continue
        wb = openpyxl.load_workbook(xlsx_path, read_only=True)
        result[m.group(1)] = list(wb.sheetnames)
        wb.close()
    return result


# ─── Property 8: 审定表块科目码属于本循环 ─────────────────────────────────────


def test_property_8_adjudication_account_codes(g_mappings):
    """审定表块的 account_codes 应包含本循环科目。"""
    adj_blocks = [m for m in g_mappings if "审定表" in m.get("sheet", "")]
    assert len(adj_blocks) >= 14, f"审定表块不足 14 个（实际 {len(adj_blocks)}）"

    for block in adj_blocks:
        wp_code = block["wp_code"]
        expected_code = G_ACCOUNT_CODES.get(wp_code)
        if not expected_code:
            continue
        codes = block.get("account_codes", [])
        if not codes:
            # 空占位块（4.1 之前的遗留），跳过
            continue
        assert expected_code in codes, (
            f"{wp_code} 审定表块 account_codes {codes} 不含 {expected_code}"
        )


# ─── Property 9: 损益类口径 = 本期发生额 ─────────────────────────────────────


def test_property_9_pl_cycles_use_current_period(g_mappings):
    """损益类循环的公式应使用本期发生额而非期末余额。"""
    for block in g_mappings:
        wp_code = block["wp_code"]
        if wp_code not in PL_CYCLES:
            continue
        cells = block.get("cells", [])
        for cell in cells:
            formula = cell.get("formula", "")
            if not formula or "TB(" not in formula:
                continue
            # 损益类 TB 公式应含「本期发生额」，不应含「期末余额」
            if "期末余额" in formula:
                # 🔴 排除 PREV() 里嵌套的 TB 引用（PREV 取上期的期末是对的）
                if "PREV(" in formula:
                    continue
                pytest.fail(
                    f"{wp_code}|{block.get('sheet', '')} 的 {cell.get('cell_ref', '')}"
                    f" 使用「期末余额」（损益类应用「本期发生额」）: {formula}"
                )


# ─── Property 10: 披露块覆盖完整 ─────────────────────────────────────────


def test_property_10_disclosure_blocks_complete(g_mappings):
    """每个有披露 sheet 的 G 循环都有对应的两变体披露块。"""
    disc_blocks = [
        m for m in g_mappings
        if "附注" in m.get("sheet", "") or "披露" in m.get("sheet", "")
    ]
    # 按 wp_code 分组
    by_code = {}
    for b in disc_blocks:
        by_code.setdefault(b["wp_code"], []).append(b)

    # G0 无披露 sheet，跳过；其余 G1~G14 应各有 2 个变体
    for code in ["G1", "G2", "G3", "G4", "G5", "G6", "G7", "G8", "G9",
                 "G10", "G11", "G12", "G13", "G14"]:
        blocks = by_code.get(code, [])
        assert len(blocks) >= 2, (
            f"{code} 只有 {len(blocks)} 个披露块（应至少 2 个：上市+国企）"
        )


def test_disclosure_presets_script_check():
    """幂等脚本 --check 应返回 0 项欠账。

    🔴 `encoding="utf-8"` 必须显式声明：Windows 上 ``text=True`` 用 locale 编码
    （GBK）解码子进程输出，而脚本输出含中文 → reader thread 抛
    ``UnicodeDecodeError`` → ``result.stdout`` 变 ``None`` → 断言以
    ``TypeError: argument of type 'NoneType' is not iterable`` 恒失败。
    表现为「守卫恒红且零信号」——既判不出脚本真有欠账，也判不出脚本是好的。
    """
    result = subprocess.run(
        [sys.executable, str(DISCLOSURE_SCRIPT), "--check"],
        capture_output=True,
        text=True,
        encoding="utf-8",
        errors="replace",
        cwd=str(BACKEND),
    )
    assert result.stdout is not None, (
        "子进程 stdout 解码失败（应已由 encoding='utf-8' 兜住）"
    )
    assert result.returncode == 0, f"--check failed:\n{result.stdout}\n{result.stderr}"
    assert "0 项欠账" in result.stdout, f"--check 输出未见「0 项欠账」:\n{result.stdout}"


# ─── 反向自检：已纠偏的错码不得复活 ─────────────────────────────────────────


KNOWN_BAD_CODES = {
    "G4": ["1501"],  # 旧值是持有至到期投资
    "G6": ["1505", "1503"],  # 旧值是债权投资减值准备 / 可供出售
    "G8": ["1506", "1503"],  # 旧值是其他债权投资 / 可供出售
    "G9": ["1507", "1510", "1504"],  # 旧值连续偏移
    "G14": ["6701"],  # 旧值是资产减值损失（IS-016↔017 互换）
    "G12": ["6115"],  # 旧值是资产处置损益
}


def test_known_bad_codes_not_present(g_mappings):
    """已实证的错码不得复活。"""
    for block in g_mappings:
        wp_code = block["wp_code"]
        bad_set = KNOWN_BAD_CODES.get(wp_code)
        if not bad_set:
            continue
        codes = block.get("account_codes", [])
        for cell in block.get("cells", []):
            formula = cell.get("formula", "")
            for bad in bad_set:
                if bad in codes:
                    pytest.fail(
                        f"{wp_code} account_codes 含已纠偏的错码 {bad}"
                    )
                if f"TB('{bad}'" in formula or f"TB(\"{bad}\"" in formula:
                    pytest.fail(
                        f"{wp_code}|{block.get('sheet', '')} 公式含已纠偏错码"
                        f" {bad}: {formula[:80]}"
                    )


# ─── Property 11: 块 sheet ∈ 源 xlsx 真实 tab（2026-09-27 补，GC-8）──────────

#: 曾经「贴错标签」的 sheet 名 → 源 xlsx 真实 tab 名（**历史台账，非逃逸阀**）。
#:
#: 形态照 `test_i_cycle_formula_presets._HISTORICAL_SHEET_LABEL_FIXES`：
#: 台账**不做豁免**，只做两件事 ①记载历史 ②由 :func:`test_no_stale_sheet_label_ledger`
#: 做 stale 检测 —— 表里的错名一旦在数据里重新出现即打红。
#:
#: 🔴 2026-09-27（spec `g-cycle-sync-foundation-and-first-canary` Task 7）：
#: 两条错名都照抄了 G11 的「明细分析表」前缀，而 G11 的 `明细分析表G11-2` **是真名**。
#: 由 `fix_g_cycle_prefill_presets.RENAME_SHEETS` 改名（**不是删块** ——
#: sheet 存在只是名字写错，删块会丢掉 4+3=7 条有效预设）。
_HISTORICAL_SHEET_LABEL_FIXES: dict[tuple[str, str], str] = {
    ("G13", "明细分析表G13-2"): "明细表G13-2",
    ("G14", "明细分析表G14-2"): "明细表G14-2",
}


def test_property_11_every_block_sheet_exists_in_template(g_mappings, template_sheets):
    """每个 G 块的 `sheet` 必须存在于对应源模板 xlsx 的 `sheetnames` 中。

    🔴 **不 strip** —— G 目录确实有名字带空格的 tab
    （`交易性金融资产实质性程序表G1A ` 尾部空格 / `信用减值损失审计程序表G14A -修订前` 名中空格），
    strip 后比较会把「空格错位」这一类缺陷整类放过。

    🔴 **无白名单** —— 源模板是唯一裁决者，预设必须对齐它。
    """
    bad: list[str] = []
    for blk in g_mappings:
        wp, sheet = blk.get("wp_code"), blk.get("sheet")
        sheetnames = template_sheets.get(wp)
        if sheetnames is None:
            # G0 走 confirmation 组件族、无 G0 单册？—— 实测有 `G0 投资循环函证.xlsx`，
            # 故这里落进来只可能是模板目录缺文件，属实缺陷不放过。
            bad.append(f"{wp}: 源模板目录下找不到 {wp} 对应的 xlsx")
            continue
        if sheet in sheetnames:
            continue
        hint = ""
        if (wp, sheet) in _HISTORICAL_SHEET_LABEL_FIXES:
            hint = (
                f"\n    🔴 该 sheet 名曾于 2026-09-27 被修正为 "
                f"'{_HISTORICAL_SHEET_LABEL_FIXES[(wp, sheet)]}'，本次出现说明数据被回退"
                "（`prefill_formula_mapping.json` 是多 spec 共享的回退高发文件）⇒ 重跑 "
                "`python backend/scripts/fix/fix_g_cycle_prefill_presets.py`"
            )
        near = [n for n in sheetnames if sheet and len(sheet) > 4 and sheet[-5:] in n]
        bad.append(
            f"{wp} 的 sheet {sheet!r} 不存在于模板中。形近真名: {near}{hint}"
        )
    assert not bad, (
        "以下 G 块指向源模板不存在的 sheet（预设在公式管理页看得见，但预填时写不进）:\n  "
        + "\n  ".join(bad)
    )


def test_no_stale_sheet_label_ledger(template_sheets):
    """历史台账里的每个「错名」都必须**确实是错的**、每个「真名」确实存在（stale 检测）。

    防两种退化：台账变逃逸阀（有人把真名塞进来当豁免）/ 台账过期（源模板真加了该 tab）。
    """
    problems: list[str] = []
    for (wp, wrong), right in _HISTORICAL_SHEET_LABEL_FIXES.items():
        sheetnames = template_sheets.get(wp, [])
        if wrong in sheetnames:
            problems.append(
                f"台账过期：{wrong!r} 现在**确实存在**于 {wp} 源模板中 ⇒ 它不再是错名，"
                "请从 _HISTORICAL_SHEET_LABEL_FIXES 移除该条目"
            )
        if right not in sheetnames:
            problems.append(
                f"台账的「正确名」{right!r} 不在 {wp} 源模板中（{sheetnames}）⇒ 台账本身写错了"
            )
    assert not problems, "sheet 名台账自检失败:\n  " + "\n  ".join(problems)


def test_property_11_reverse_self_check_would_catch_a_ghost_sheet(template_sheets):
    """反向自检：判据确实能识别幽灵 sheet（防判据空转）。

    合成一个不存在的名字塞进比对逻辑，必须被认定为「不在模板里」。
    """
    assert template_sheets, "模板 sheetnames 为空 —— 判据无基础"
    ghost = "明细分析表G13-2"  # 已修正的历史错名，现在确实不该存在
    assert ghost not in template_sheets["G13"], (
        f"{ghost!r} 居然在 G13 模板里 —— 反向自检的前提不成立"
    )
    assert "明细表G13-2" in template_sheets["G13"]


def test_prefill_script_check_is_clean():
    """`fix_g_cycle_prefill_presets.py --check` 应 0 欠账（幂等锚点）。

    🔴 `encoding="utf-8"` 必须显式声明（同本文件既有 `test_disclosure_presets_script_check`
    的理由：Windows 下 locale GBK 解码中文 stdout 会让 `stdout` 变 None，
    断言以 TypeError 恒失败 = 判据恒红且零信号）。
    """
    script = BACKEND / "scripts" / "fix" / "fix_g_cycle_prefill_presets.py"
    result = subprocess.run(
        [sys.executable, str(script), "--check"],
        capture_output=True,
        text=True,
        encoding="utf-8",
        errors="replace",
        cwd=str(BACKEND.parent),
    )
    assert result.stdout is not None, "子进程 stdout 解码失败（应已由 encoding='utf-8' 兜住）"
    assert result.returncode == 0, f"--check 有欠账:\n{result.stdout}\n{result.stderr}"

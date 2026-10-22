"""合并附注种子多级表头与现行四版 Word 一致（spec consol-note-header-library-and-node-reaggregate R6.3 / design §15.3）。

误修脚本 ``_fix_missing_multi_header.py``（已删）从平表头反推多级表头；Word 生成器的表头行数判定又把带合并单元格的数据行
算进表头。一次性脚本按 Word 修了 30 张表、统一了全部多级表头表的平表头。本测试在测试时解析 Word（``note_word_table``：
表头行数、表头空白格归一；标题行库 / Word 同步工具 / 上年附注导入同一套）逐表钉住：

- 跨度错的 17 张：每列表头路径（展示 / Word 导出 / 标题行库同一 ``header_rows`` 投影）= Word 每列表头路径，表头行数 = Word。
- 数据行（含空白占位行）被误升进表头的 12 张：表头行数按 Word 网格人工核定；表头路径 = Word（空白格归一后）；
  表体 = Word 数据行（网格位置；误修重复插入的行已去掉，Word 表后面嵌着的另一张表不算）。
- 人工定的 1 张（上市 六-1-2）：Word 表头空白格左右都是分组，平台规则按独立列，种子按语义归「本期减少」。
- 全部多级表头表：平表头 = 展示表头逐列全路径（公式管理列名、校验规则名用它）。
- 表头空白格归一规则（``_normalize_header_gaps``）按四版 Word 里的全部出现处逐条钉住。
"""
from __future__ import annotations

import json
import re
from functools import lru_cache
from pathlib import Path

import pytest
from docx import Document
from docx.oxml.ns import qn

from app.services.consol_note_formula_service import multi_header_to_column_groups
from app.services.consol_note_headers import header_cells, multi_header_rows, normalize_consol_note_section
from app.services.note_table_header_library import column_paths, norm
from app.services.note_word_table import (
    _header_row_count, _logical_columns, _normalize_header_gaps, _table_grid, parse_word_table, seed_multi_header,
)

DATA = Path(__file__).resolve().parents[1] / "data"
DOCX = DATA / "audit_report_templates" / "disclosure_notes"
_TAG = re.compile(r"\{\{table:([^}]+)\}\}")
_STYLE_REF = re.compile(r"##STYLE_REF[^#]*##")

#: 跨度错：(口径, section_id) → Word 表标记；None = Word 表前没有标记，按表前标题定位
SPAN_FIXED: dict[tuple[str, str], str | None] = {
    ("soe", "五-5-1"): "八、5:0", ("soe", "五-9-6"): "八、9:5", ("soe", "五-13-2"): "八、12:1",
    ("soe", "五-18-1"): "八、17:0", ("soe", "五-19-2"): "八、18:1", ("soe", "五-19-9"): "八、18:8",
    ("soe", "五-19-10"): "八、18:9", ("soe", "五-22-1"): "八、21:0", ("soe", "五-22-2"): "八、21:1",
    ("soe", "五-29-1"): "八、28:0", ("soe", "五-55-8"): "八、54:7",
    ("listed", "五-10-8"): "五、10:5", ("listed", "五-11-3"): "五、11:1", ("listed", "五-18-1"): "五、18:0",
    ("listed", "五-24-2"): None, ("listed", "五-78-9"): None, ("listed", "十六-3"): "十六:2",
}
#: 数据行被误升：(口径, section_id) → (Word 表标记, 表头行数, 表体行数)
ROW_FIXED: dict[tuple[str, str], tuple[str, int, int]] = {
    ("soe", "五-57-2"): ("八、56:1", 1, 4), ("soe", "五-64-4"): ("八、64:3", 3, 9),
    ("listed", "九-1-27"): ("九、1:26", 1, 4), ("listed", "九-1-28"): ("九、1:27", 1, 4),
    ("listed", "五-9-4"): ("五、9:3", 3, 4), ("listed", "五-9-5"): ("五、9:4", 3, 4),
    ("listed", "五-9-6"): ("五、9:5", 3, 4), ("listed", "五-9-7"): ("五、9:6", 3, 4),
    ("listed", "五-10-1"): ("五、10:0", 2, 9), ("listed", "五-11-1"): ("五、11:0", 2, 9),
    ("listed", "五-62-4"): ("五、62:3", 3, 9), ("listed", "三-重要会计政策会-1"): ("三、重要会计政策、会:0", 1, 2),
}


def _clean(text) -> str:
    return str(text or "").replace("\n", "").strip()


@lru_cache(maxsize=None)
def _seed(std: str) -> dict[str, dict]:
    raw = (DATA / f"consol_note_sections_{std}.json").read_bytes().decode("utf-8")
    return {s["section_id"]: s for s in json.loads(raw)}


@lru_cache(maxsize=None)
def _word_tables(std: str, variant: str = "consolidated") -> tuple[tuple[str | None, str, object], ...]:
    """Word 每张表 (表前标记, 表前标题, 表元素)。标记只管紧跟的一张表（``sync_note_templates_from_word`` 同口径）。"""
    doc = Document(str(DOCX / f"{std}_{variant}.docx"))
    out: list[tuple[str | None, str, object]] = []
    tag, title = None, ""
    for el in doc.element.body:
        if el.tag == qn("w:p"):
            text = "".join(t.text or "" for t in el.iter(qn("w:t"))).strip()
            hit = _TAG.search(text)
            if hit:
                tag = hit.group(1)
                rest = _STYLE_REF.sub("", _TAG.sub("", text)).strip()
                title = rest or title
            elif text and not text.startswith("##"):
                title = text
        elif el.tag == qn("w:tbl"):
            out.append((tag, title, el))
            tag, title = None, ""
    return tuple(out)


def _locate(std: str, tag: str | None, title: str | None = None):
    tables = _word_tables(std)
    if tag is not None:
        hits = [el for t, _title, el in tables if t == tag]
    else:
        hits = [el for t, ti, el in tables if t is None and norm(ti) == norm(title)]
    assert len(hits) == 1, f"{std}：Word 表 {tag or title} 命中 {len(hits)} 张"
    return hits[0]


def _paths(grid: dict, texts: dict, n_hdr: int, cols: list[int]) -> list[tuple[str, ...]]:
    """每列自上而下的非空表头文字（同一单元格只计一次；``parse_word_table`` 同口径），规范化。"""
    out = []
    for c in cols:
        path, seen = [], set()
        for r in range(n_hdr):
            owner = grid.get((r, c))
            if owner is None or owner in seen:
                continue
            seen.add(owner)
            if _clean(texts.get(owner)):
                path.append(norm(texts[owner]))
        out.append(tuple(path))
    return out


def _full_headers(sec: dict) -> list[str]:
    width = len(sec["headers"])
    paths: list[list[str]] = [[] for _ in range(width)]
    for _r, c0, cell in sorted(header_cells(multi_header_rows(sec["multi_header"], width), width), key=lambda x: x[:2]):
        if _clean(cell["text"]):
            for c in range(c0, c0 + cell["colspan"]):
                paths[c].append(_clean(cell["text"]))
    return ["/".join(p) if p else h for p, h in zip(paths, sec["headers"])]


@pytest.mark.parametrize(("std", "sid"), sorted(SPAN_FIXED))
def test_span_fixed_tables_follow_word(std: str, sid: str) -> None:
    sec = _seed(std)[sid]
    tbl = _locate(std, SPAN_FIXED[(std, sid)], sec.get("title"))
    word = parse_word_table(tbl, "")
    assert word.columns == len(sec["headers"])
    assert len(sec["multi_header"]) == word.header_rows
    assert column_paths(sec) == [tuple(norm(x) for x in p) for p in word.header_paths]
    # 逐格 = 平台种子格式（Word 同步工具同一函数）：Word 同步重跑不会把它改回去
    assert sec["multi_header"] == seed_multi_header(tbl)[1]
    assert sec["_column_groups"] == multi_header_to_column_groups(sec["multi_header"])


@pytest.mark.parametrize(("std", "sid"), sorted(ROW_FIXED))
def test_row_fixed_tables_follow_word(std: str, sid: str) -> None:
    tag, n_hdr, n_body = ROW_FIXED[(std, sid)]
    sec = _seed(std)[sid]
    tbl = _locate(std, tag)
    nrows, ncols, grid, texts, spans = _table_grid(tbl)
    # 人工核定的表头行数 = 平台判定（数据行里的空白合并格 / 通栏分类行 / 空白占位行不算表头）
    assert min(_header_row_count(nrows, ncols, grid, texts, spans), nrows) == n_hdr
    width = len(sec["headers"])
    cols = list(range(ncols)) if width == ncols else _logical_columns(ncols, n_hdr, grid)
    assert len(cols) == width
    body = [[_clean(texts.get((r, c))) if grid.get((r, c)) == (r, c) else "" for c in cols] for r in range(n_hdr, nrows)]
    assert sec["rows"] == body[:n_body]
    grid, texts, spans = _normalize_header_gaps(ncols, n_hdr, grid, texts, spans)
    assert column_paths(sec) == _paths(grid, texts, n_hdr, cols)
    if n_hdr == 1:
        assert sec.get("multi_header") is None and "_column_groups" not in sec
        assert sec["headers"] == [_clean(texts.get((0, c))) if grid.get((0, c)) == (0, c) else "" for c in cols]
    else:
        assert sec["multi_header"] == seed_multi_header(tbl, cols)[1]
        assert sec["_column_groups"] == multi_header_to_column_groups(sec["multi_header"])


def test_manual_development_expenditure_header() -> None:
    """上市 六-1-2 开发支出：Word 第 0 行第 4 格空白未合并，左右「本期增加」「本期减少」都是分组 ⇒ 平台按独立列；
    种子按语义归「本期减少」（确认为无形资产减少开发支出余额；国企同表、上市单体同表都在「本期减少」下）。"""
    from app.services.note_table_header_library import rollforward_columns

    sec = _seed("listed")["六-1-2"]
    _n, platform = seed_multi_header(_locate("listed", "六、1:1"))
    assert (platform[0][4], platform[1][4]) == ("确认为无形资产", ""), "Word 模板已改 ⇒ 按 Word 原样核对、删掉本条"
    assert sec["multi_header"] == [
        ["项  目", "期初余额", "本期增加", "", "本期减少", "", "", "期末余额"],
        ["", "", "内部开发支出", "其他增加", "确认为无形资产", "计入当期损益", "其他减少", ""],
    ]
    assert rollforward_columns(sec).decrease == (4, 5, 6)


@pytest.mark.parametrize("std", ["soe", "listed"])
def test_flat_headers_are_full_display_paths(std: str) -> None:
    """多级表头表的平表头 = 展示表头逐列全路径（Word 留白的标签列保留原占位名）。"""
    bad = [
        sid for sid, s in _seed(std).items()
        if isinstance(s.get("multi_header"), list) and len(s["multi_header"]) >= 2 and s["headers"] != _full_headers(s)
    ]
    assert bad == []


@pytest.mark.parametrize("std", ["soe", "listed"])
def test_no_blank_header_row_and_every_table_projects(std: str) -> None:
    """整行全空的表头行只能是被误升的空数据行；每张表的表头跨度都能投影（覆盖每一列、不重叠）。"""
    blank = [sid for sid, s in _seed(std).items()
             if any(not any(_clean(c) for c in row) for row in s.get("multi_header") or [])]
    assert blank == []
    for sec in _seed(std).values():
        normalize_consol_note_section(sec)


def test_bad_multi_header_fixer_is_gone() -> None:
    """误修脚本已删：它从平表头反推多级表头，重跑会把这些表改回去。"""
    assert not (DATA.parent / "scripts" / "seed" / "_fix_missing_multi_header.py").exists()


def _word_grid_paths(std: str, tag: str, variant: str = "consolidated") -> tuple[int, list[tuple[str, ...]]]:
    tables = _word_tables(std, variant)
    hits = [el for t, _title, el in tables if t == tag]
    assert len(hits) == 1, (std, variant, tag, len(hits))
    nrows, ncols, grid, texts, spans = _table_grid(hits[0])
    n_hdr = min(_header_row_count(nrows, ncols, grid, texts, spans), nrows)
    grid, texts, spans = _normalize_header_gaps(ncols, n_hdr, grid, texts, spans)
    return n_hdr, _paths(grid, texts, n_hdr, list(range(ncols)))


class TestWordHeaderRules:
    """``note_word_table`` 的表头行数判定与表头空白格归一，按四版 Word 里的全部出现处钉住（两版同表都核）。"""

    @pytest.mark.parametrize(("std", "tag", "n_hdr"), [
        ("soe", "八、5:0", 2),            # 两级表头「账龄」列下格留白（只看纵向合并会判成一行）
        ("soe", "八、64:3", 3),           # 数据行两格一并的空白取数格不是表头
        ("listed", "五、9:5", 3), ("listed", "五、10:0", 2), ("listed", "五、11:0", 2), ("listed", "五、62:3", 3),
        ("listed", "九、1:26", 1),        # 通栏「商品价格风险」分类行
        ("listed", "三、重要会计政策、会:0", 1),  # 空白占位行
    ])
    @pytest.mark.parametrize("variant", ["consolidated", "standalone"])
    def test_header_row_count(self, std: str, tag: str, n_hdr: int, variant: str) -> None:
        tables = _word_tables(std, variant)
        hits = [el for t, _title, el in tables if t == tag]
        if variant == "standalone" and not hits:
            pytest.skip("单体版无同标记表")
        nrows, ncols, grid, texts, spans = _table_grid(hits[0])
        assert min(_header_row_count(nrows, ncols, grid, texts, spans), nrows) == n_hdr

    def test_blank_merged_cell_never_extends_header(self) -> None:
        """空白的合并格不撑开表头。四版 Word 里这类行都另被截断规则认作数据行（只去掉撑开这一侧的判据，全部 1366
        张表结果不变）；这里钉住截断规则认不出的行：首列留白（不按「首列有字」截断）、本行又有字（不是空白占位行）。"""
        t = Document().add_table(rows=3, cols=4)
        t.cell(1, 1).merge(t.cell(1, 2))
        for (r, c), text in {(0, 0): "项目", (0, 1): "期末余额", (0, 2): "期初余额", (0, 3): "备注",
                             (1, 3): "说明", (2, 0): "合计", (2, 1): "1", (2, 2): "2"}.items():
            t.cell(r, c).text = text
        nrows, ncols, grid, texts, spans = _table_grid(t._tbl)
        assert spans[(1, 1)] == (1, 2) and texts[(1, 1)] == ""
        assert _header_row_count(nrows, ncols, grid, texts, spans) == 1

    @pytest.mark.parametrize(("std", "tag", "col", "path"), [
        # 1. 子表头从空白格跨进右邻分组 ⇒ 空白格并入右邻（存货按组合计提跌价准备四张表）
        ("listed", "五、9:3", 1, ("期末余额", "账面余额", "金额")), ("listed", "五、9:4", 1, ("上年年末余额", "账面余额", "金额")),
        ("listed", "五、9:5", 1, ("期末余额", "账面余额", "金额")), ("listed", "五、9:6", 1, ("上年年末余额", "账面余额", "金额")),
        # 2a. 左邻是分组、右邻不是 ⇒ 列名属左邻分组
        ("soe", "八、4:5", 5, ("本期变动情况", "其他变动")),
        ("listed", "五、15:3", 5, ("期末余额", "逾期本金")), ("listed", "五、15:4", 5, ("上年年末余额", "逾期本金")),
        # 2b. 左邻是单列 ⇒ 独立一列
        ("soe", "八、11:1", 7, ("原因",)),
        ("listed", "五、18:0", 2, ("减值准备期初余额",)), ("listed", "十六:2", 2, ("减值准备期初余额",)),
        # 2c. 左右都是分组 ⇒ 分不出，按独立列（种子人工定，见 test_manual_development_expenditure_header）
        ("listed", "六、1:1", 4, ("确认为无形资产",)),
        # 2d. 左邻是分组但属另一个上级（第 1 行：「以摊余成本计量」属「期末余额」）⇒ 不并，独立一列
        ("listed", "九、1:2", 6, ("上年年末余额", "以公允价值计量且其变动计入当期损益")),
        ("listed", "九、1:2", 1, ("期末余额", "以公允价值计量且其变动计入当期损益")),
    ])
    @pytest.mark.parametrize("variant", ["consolidated", "standalone"])
    def test_header_gap_normalization(self, std: str, tag: str, col: int, path: tuple, variant: str) -> None:
        if variant == "standalone" and not any(t == tag for t, _title, _el in _word_tables(std, variant)):
            pytest.skip("单体版无同标记表")
        _n, paths = _word_grid_paths(std, tag, variant)
        assert paths[col] == tuple(norm(x) for x in path)

    def test_gaps_only_where_listed(self) -> None:
        """归一只动上面这些表：四版 Word 里其余表的表头列路径与不归一时相同。"""
        listed = {("soe", "八、4:5"), ("soe", "八、11:1"), ("listed", "五、15:3"), ("listed", "五、15:4"),
                  ("listed", "五、18:0"), ("listed", "十六:2"), ("listed", "六、1:1"), ("listed", "九、1:2"),
                  *(("listed", f"五、9:{i}") for i in range(3, 7))}
        touched = set()
        for std in ("soe", "listed"):
            for variant in ("consolidated", "standalone"):
                for tag, _title, el in _word_tables(std, variant):
                    nrows, ncols, grid, texts, spans = _table_grid(el)
                    if not nrows:
                        continue
                    n_hdr = min(_header_row_count(nrows, ncols, grid, texts, spans), nrows)
                    g2, t2, _s2 = _normalize_header_gaps(ncols, n_hdr, grid, texts, spans)
                    if (g2, t2) != (grid, texts):
                        touched.add((std, tag))
        assert touched <= listed and ("listed", "五、9:3") in touched


def test_word_sync_tool_reads_the_same_headers() -> None:
    """Word 同步工具（长期可重跑）表头走平台同一套：修过的表重跑同步，多级表头与种子逐格相同（不会改回去）。"""
    from scripts.seed.sync_note_templates_from_word import extract_all_word_tables

    by_tag = {wt["tag"]: wt for wt in extract_all_word_tables(DOCX / "listed_consolidated.docx") if wt["tag"]}
    seed = _seed("listed")
    for sid, tag in (("五-9-4", "五、9:3"), ("五-9-6", "五、9:5"), ("五-10-1", "五、10:0"), ("五-11-1", "五、11:0"),
                     ("五-62-4", "五、62:3"), ("五-18-1", "五、18:0"), ("五-11-3", "五、11:1")):
        wt = by_tag[tag]
        assert wt["n_header_rows"] == len(seed[sid]["multi_header"]), sid
        assert wt["multi_header"] == seed[sid]["multi_header"], sid
    assert by_tag["九、1:26"]["n_header_rows"] == 1 and by_tag["九、1:26"]["multi_header"] is None

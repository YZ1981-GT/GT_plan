"""审计报告正文模板后处理（spec deliverable-trio-generation-accuracy §3，report_body_postprocess）。

夹具用 ``backend/data/audit_report_templates/report_body`` 下的真实致同模板（2026-10-16 普查 34 份）：
- 上市 / 三板模板 OPT 开闭标记错配（「##OPT:key_audit_matters##」配「##/OPT:other_information##」）；
- 单体非标模板「##NOTE:单体复核:…##」未闭合、证监会指引单格表、「【非上市实体删除（项目合伙人）】」；
- 落款日期写死「二〇二六年 X 月XX日」/「二○二二年X月XX日」；全部模板写「股东权益变动表」。
"""

from __future__ import annotations

import threading
from datetime import date
from pathlib import Path

import pytest
from docx import Document
from docx.oxml.ns import qn

import app.services.report_body_postprocess as rbp
from app.services.report_body_section_blocks import scan_report_body_sections
from app.services.template_manifest_loader import resolve_template_base_dir
from app.services.word_doc_utils import (
    apply_optional_sections,
    replace_placeholders_in_doc,
    scan_optional_sections,
    strip_guidance_notes,
)

BASE = resolve_template_base_dir() / "report_body"
MAPPING = {
    "company_full_name": "重药控股安徽有限公司",
    "company_short_name": "重药安徽公司",
    "responsibility_organ": "董事会",
    "audit_period_end": "2025年12月31日",
    "audit_year": "2025",
    "prior_year": "2024",
    "report_number": "[报告文号]",
    "report_date": "[报告日期]",
    "report_scope": "",
}


def _tpl(pattern: str) -> Path:
    hits = sorted(BASE.glob(pattern))
    if not hits:
        pytest.skip(f"报告正文模板缺失：{pattern}")
    return hits[0]


def _doc(pattern: str, mapping: dict | None = None):
    doc = Document(str(_tpl(pattern)))
    replace_placeholders_in_doc(doc, mapping or MAPPING)
    return doc


def _text(doc) -> str:
    return "\n".join("".join(t.text or "" for t in p.iter(qn("w:t"))) for p in doc.element.body.iter(qn("w:p")))


def _heads(doc) -> list[str]:
    return [f"{num}、{name}" for _el, num, name in rbp._headings(doc)]


# ── OPT 标记按章节标题重放 ────────────────────────────────────────────────────


def _opt_templates() -> list[Path]:
    out = []
    for f in sorted(BASE.rglob("*.docx")):
        if any(p.text.strip().startswith(("##OPT:", "##/OPT:")) for p in Document(str(f)).paragraphs):
            out.append(f)
    return out


def test_mismatched_opt_markers_are_the_defect_being_fixed():
    """前提：上市简版无保留模板的开闭标记 id 不配对，旧扫描一个可选段也认不出（判据不空转）。"""
    doc = _doc("standalone/1.1 模板A-无保留*.docx")
    assert scan_optional_sections(doc) == []
    assert rbp.normalize_optional_markers(doc) == 2
    assert {s.section_id for s in scan_optional_sections(doc)} == {"key_audit_matters", "other_information"}


@pytest.mark.parametrize("path", _opt_templates(), ids=lambda p: p.name[:16])
def test_every_opt_block_wraps_exactly_its_own_section(path: Path):
    doc = Document(str(path))
    rbp.normalize_optional_markers(doc)
    first = [(s.section_id, s.start_index, s.end_index) for s in scan_optional_sections(doc)]
    assert first, "带 OPT 标记的模板重放后应识别出可选段"
    paras = doc.paragraphs
    for sec in scan_optional_sections(doc):
        inner = [p.text.strip() for p in paras[sec.start_index + 1: sec.end_index - 1]]
        heads = [t for t in inner if rbp._HEADING_RE.match(t) and len(t) <= 40]
        assert len(heads) == 1, (sec.section_id, heads)  # 块内恰好是本段一个一级标题
        assert any(t in heads[0] for t in rbp.OPT_SECTION_TITLES[sec.section_id])
        nxt = paras[sec.end_index].text.strip() if sec.end_index < len(paras) else ""
        assert rbp._HEADING_RE.match(nxt) or nxt.startswith("##OPT:"), (sec.section_id, nxt)
    rbp.normalize_optional_markers(doc)  # 幂等
    assert [(s.section_id, s.start_index, s.end_index) for s in scan_optional_sections(doc)] == first


def test_deselected_kam_leaves_no_marker_and_headings_close_up():
    doc = _doc("standalone/1.1 模板A-无保留*.docx")
    rbp.prepare_report_body(doc, opinion_choice="unqualified", audit_year=2025, company_name=MAPPING["company_full_name"])
    apply_optional_sections(doc, {"key_audit_matters": False, "other_information": True})
    strip_guidance_notes(doc)
    rbp.finalize_report_body(doc, standalone=True, listed=True)
    assert "##" not in _text(doc)
    assert _heads(doc)[:4] == ["一、审计意见", "二、形成审计意见的基础", "三、其他信息", "四、管理层和治理层对财务报表的责任"]
    assert not any("关键审计事项" in h for h in _heads(doc))


# ── 强调事项段 ────────────────────────────────────────────────────────────────


def test_emphasis_section_inserted_after_basis_and_numbering_follows():
    doc = _doc("standalone/3.1 模板C-无保留*.docx")
    res = rbp.prepare_report_body(doc, opinion_choice="unqualified_with_emphasis", audit_year=2025,
                                  company_name=MAPPING["company_full_name"])
    assert res.emphasis_inserted
    assert _heads(doc) == ["一、审计意见", "二、形成审计意见的基础", "三、强调事项",
                           "四、管理层和治理层对财务报表的责任", "五、注册会计师对财务报表审计的责任"]
    heads = {el: name for el, _n, name in rbp._headings(doc)}
    emph = next(el for el, name in heads.items() if name == "强调事项")
    basis = next(el for el, name in heads.items() if name == "形成审计意见的基础")
    assert rbp._ptext(emph.getnext()) == rbp.EMPHASIS_TEXT
    # 标题沿用基础段标题格式（加粗），章节锚点能识别出强调事项段
    assert (emph.find(".//" + qn("w:b")) is not None) == (basis.find(".//" + qn("w:b")) is not None) is True
    assert "emphasis" in [s.section_id for s in scan_report_body_sections(doc)]
    assert "强调事项段（所强调的事项与附注编号）" in rbp.pending_items(doc)
    assert rbp.insert_emphasis_section(doc) is False  # 已有则不重复插


def test_standard_unqualified_has_no_emphasis_section():
    doc = _doc("standalone/3.1 模板C-无保留*.docx")
    rbp.prepare_report_body(doc, opinion_choice="unqualified", audit_year=2025, company_name="某有限公司")
    assert not any("强调事项" in h for h in _heads(doc))


def test_emphasis_stays_outside_kam_optional_block():
    """上市模板：基础段之后紧跟 KAM 可选段，强调事项插在 KAM 开标记之前 —— 不勾选 KAM 时强调事项不能被一并删掉。"""
    doc = _doc("standalone/1.1 模板A-无保留*.docx")
    rbp.prepare_report_body(doc, opinion_choice="unqualified_with_emphasis", audit_year=2025, company_name="某股份有限公司")
    apply_optional_sections(doc, {"key_audit_matters": False, "other_information": False})
    strip_guidance_notes(doc)
    rbp.finalize_report_body(doc, standalone=True, listed=True)
    assert _heads(doc)[:4] == ["一、审计意见", "二、形成审计意见的基础", "三、强调事项", "四、管理层和治理层对财务报表的责任"]


# ── 编制提示 / 单体口径 ───────────────────────────────────────────────────────


def _finalized(pattern: str, *, standalone: bool, listed: bool, opinion: str = "qualified", company: str = "重药控股安徽有限公司"):
    doc = _doc(pattern)
    rbp.prepare_report_body(doc, opinion_choice=opinion, audit_year=2025, company_name=company)
    apply_optional_sections(doc, {s.section_id: True for s in scan_optional_sections(doc)})
    strip_guidance_notes(doc)
    st = rbp.finalize_report_body(doc, standalone=standalone, listed=listed)
    return doc, st


def test_standalone_qualified_report_has_no_guidance_or_group_clause():
    doc, st = _finalized("standalone/3. 模板C-保留*.docx", standalone=True, listed=False)
    text = _text(doc)
    assert st.guide_boxes == 1 and st.group_clauses == 1
    for leftover in ("##", "监管规则适用指引", "【", "集团审计", "合并资产负债表", "合并财务报表"):
        assert leftover not in text, leftover
    # 示例段只去标记、措辞改单体，段落本身保留（须由审计师按实际事项改写）
    assert "资产负债表上反映的账面价值" in text
    assert _heads(doc)[:2] == ["一、保留意见", "二、形成保留意见的基础"]
    assert "所有者权益变动表" in text and "股东权益" not in text


def test_disclaimer_partner_label_by_listing_status_and_signing_year():
    doc, _st = _finalized("standalone/1.4.1*.docx", standalone=True, listed=False, opinion="disclaimer")
    text = _text(doc)
    assert "（项目合伙人）" not in text and "【" not in text
    assert "对财务报表发表审计意见的基础" in text  # 「对合并财务报表…」改为单体措辞
    assert "独立性要求（如适用），我们独立于" in text and "（如适用）/" not in text  # 二选一措辞已定
    assert "二〇二六年X月XX日" in text and "二○二二" not in text
    listed_doc, _ = _finalized("standalone/1.4.1*.docx", standalone=True, listed=True, opinion="disclaimer",
                               company="某某股份有限公司")
    listed_text = _text(listed_doc)
    assert "（项目合伙人）" in listed_text and "非上市实体删除" not in listed_text
    assert "股东权益变动表" in listed_text  # 股份有限公司保留「股东权益」


def test_signing_date_written_in_chinese_when_report_date_known():
    doc = _doc("standalone/3.1 模板C-无保留*.docx")
    rbp.prepare_report_body(doc, opinion_choice="unqualified", audit_year=2025, company_name="某有限公司",
                            report_date="2026-03-05")
    text = _text(doc)
    assert "二〇二六年三月五日" in text and "X" not in text.split("致同会计师事务所")[-1]
    assert "落款日期（月、日）" not in rbp.pending_items(doc)
    assert rbp.parse_report_date("[报告日期]") is None and rbp.parse_report_date("2026年3月5日") == date(2026, 3, 5)


def test_equity_statement_name_follows_company_form_including_toc():
    doc = _doc("standalone/3.1 模板C-无保留*.docx")
    assert rbp.normalize_equity_statement_name(doc, "重药控股安徽有限公司") >= 2  # 目录表 + 审计意见段
    text = _text(doc)
    assert "股东权益" not in text and text.count("所有者权益变动表") >= 2
    assert rbp.equity_statement_name("某某股份有限公司") == "股东权益变动表"
    assert rbp.equity_statement_name("某某有限责任公司") == "所有者权益变动表"
    joint = _doc("standalone/3.1 模板C-无保留*.docx")
    assert rbp.normalize_equity_statement_name(joint, "某某股份有限公司") == 0


# ── 待补充清单 ────────────────────────────────────────────────────────────────


def test_missing_fields_only_template_fields_in_chinese():
    raw = Document(str(_tpl("standalone/3.1 模板C-无保留*.docx")))
    used = rbp.placeholder_keys(raw)
    assert "report_number" in used and "signing_partner" not in used  # 落款表里签名手写，模板不用该占位
    labels = rbp.missing_field_labels(["signing_partner", "report_number", "report_date"], used)
    assert labels == ["报告文号"]


def test_pending_items_for_qualified_template():
    doc = _doc("standalone/3. 模板C-保留*.docx")
    rbp.prepare_report_body(doc, opinion_choice="qualified", audit_year=2025, company_name="某有限公司")
    items = rbp.pending_items(doc)
    assert "形成保留意见的基础段（模板示例，须按实际事项改写）" in items
    assert "落款日期（月、日）" in items
    assert any(i.startswith("正文「×」待填") for i in items)


def test_toc_page_placeholders_listed_as_pending():
    """目录表页码「1-X」「7-X」生成时算不出 → 列进待补充（Playwright 实测：交付件留着「1-X」而预览只报文号 / 日期）。"""
    doc = Document()
    tbl = doc.add_table(rows=3, cols=2)
    for row, (name, page) in zip(tbl.rows, [("审计报告", "1-X"), ("资产负债表", "1-2"), ("财务报表附注", "7－Ｘ")]):
        row.cells[0].text, row.cells[1].text = name, page
    assert rbp.toc_pending_pages(doc) == ["审计报告 1-X", "财务报表附注 7－Ｘ"]
    assert "目录页码（审计报告 1-X、财务报表附注 7－Ｘ，装订后按实际页数填写）" in rbp.pending_items(doc)
    # 正文段落里的「1-X」不是目录行，不报
    plain = Document()
    plain.add_paragraph("1-X")
    assert rbp.toc_pending_pages(plain) == []


@pytest.mark.parametrize("path", sorted(BASE.rglob("*.docx")), ids=lambda p: p.name[:16])
def test_every_template_toc_pages_reported(path: Path):
    """34 份模板的目录都有「审计报告 1-X」「财务报表附注 7-X」：生成后都要出现在待补充里。"""
    doc = Document(str(path))
    replace_placeholders_in_doc(doc, MAPPING)
    assert rbp.toc_pending_pages(doc) == ["审计报告 1-X", "财务报表附注 7-X"]
    assert any(i.startswith("目录页码（审计报告 1-X、财务报表附注 7-X") for i in rbp.pending_items(doc))


def test_strip_guidance_notes_terminates_on_labelled_inline_note():
    """``##/NOTE:label##`` 闭标记：旧正则认不出 ⇒ 同一段无限重试（进程挂死）。在线程里跑并限时。"""
    doc = Document()
    doc.add_paragraph("前文##NOTE:提示##内部指引##/NOTE:提示##后文")
    doc.add_paragraph("孤立闭标记##/NOTE##仍在")
    done = threading.Event()

    def _run():
        strip_guidance_notes(doc)
        done.set()

    threading.Thread(target=_run, daemon=True).start()
    assert done.wait(10), "strip_guidance_notes 未终止"
    texts = [p.text for p in doc.paragraphs]
    assert texts[0] == "前文后文"
    assert texts[1] == "孤立闭标记##/NOTE##仍在"  # 不成对的残留留给 strip_report_guidance 清
    rbp.strip_report_guidance(doc, standalone=False, listed=False)
    assert [p.text for p in doc.paragraphs][1] == "孤立闭标记仍在"


def test_strip_guidance_notes_terminates_when_close_precedes_open():
    """闭标记在开标记之前：正则替换不掉，修复前同样对这一段无限重试。"""
    doc = Document()
    doc.add_paragraph("倒序##/NOTE##中间##NOTE:x##")
    done = threading.Event()
    threading.Thread(target=lambda: (strip_guidance_notes(doc), done.set()), daemon=True).start()
    assert done.wait(10), "strip_guidance_notes 未终止"
    rbp.strip_report_guidance(doc, standalone=False, listed=False)
    assert doc.paragraphs[0].text == "倒序中间"


# ── 合成文档：持续经营段 / 分部标题 ─────────────────────────────────────────


def _heading_doc(*heads_and_bodies: str):
    doc = Document()
    for text in heads_and_bodies:
        p = doc.add_paragraph()
        p.add_run(text).bold = bool(rbp._HEADING_RE.match(text))
    return doc


def test_emphasis_goes_after_going_concern_section():
    doc = _heading_doc("一、审计意见", "意见段。", "二、形成审计意见的基础", "基础段。",
                       "三、与持续经营相关的重大不确定性", "持续经营段。", "四、管理层和治理层对财务报表的责任", "管理层段。")
    assert rbp.insert_emphasis_section(doc)
    rbp.renumber_section_headings(doc)
    assert _heads(doc) == ["一、审计意见", "二、形成审计意见的基础", "三、与持续经营相关的重大不确定性",
                           "四、强调事项", "五、管理层和治理层对财务报表的责任"]


def test_part_heading_keeps_its_own_number():
    doc = _doc("standalone/1.2 模版A-无保留*.docx")
    rbp.renumber_section_headings(doc)
    heads = _heads(doc)
    assert "二、按照相关法律法规的要求报告的事项" in heads
    main = [h for h in heads if "按照相关法律法规" not in h]
    assert [h.split("、", 1)[0] for h in main] == [rbp._cn_number(i) for i in range(1, len(main) + 1)]
    assert not any(h.startswith("X、") for h in heads)  # 详版「X、」占位序号已连续编号

"""D4（4）分解信息表：源 xlsx ↔ 附注模板 ↔ 前端列构造 三向比对守卫。

判据真源 = `backend/wp_templates/D/D4-1至D4-4 营业收入 - 审定表明细表（Leap-常规程序）.xlsx`
（运行时权威模板目录；`基础数据/` 下的参考副本已落后，不作判据）。

实证（openpyxl 直读）
=====================

上市 `附注披露信息（上市公司）` R44~R56 / 国企 `附注披露信息（国企）` R37~R49，
**两版逐行同构**：

- 表头三层：`B45:I45 本期发生额` > 四个类别（各 `colspan=2`）> `收入 / 成本`
- 列 **9 个**（A..I）：label + 4 类别 × 2，**没有横向合计列**
- 行 **9 个**：主营业务(SUM) / 其中：时点 / 时段 / **空可扩行** / 其他业务(SUM) /
  其中：时点 / 时段 / 租赁收入 / 合计

本文件钉死这些事实，并交叉锁死附注模板 JSON 与前端 `buildD4TransposeColumns`
（读 `.ts` 源码），防三侧任一漂移。

spec: d-cycle-four-table-extraction-and-disclosure-completion (Task 31)
Validates: Requirements 7.5
Property: 25
"""

from __future__ import annotations

import json
import re
from pathlib import Path

import pytest
from openpyxl import load_workbook


def _repo_root() -> Path:
    sentinels = (
        Path("backend") / "data" / "note_template_listed.json",
        Path("audit-platform") / "frontend" / "package.json",
    )
    cur = Path(__file__).resolve()
    for cand in [cur, *cur.parents]:
        if all((cand / s).exists() for s in sentinels):
            return cand
    raise AssertionError("未能定位仓库根（双哨兵均需存在）")


REPO_ROOT = _repo_root()
SRC_XLSX = (
    REPO_ROOT
    / "backend"
    / "wp_templates"
    / "D"
    / "D4-1至D4-4 营业收入 - 审定表明细表（Leap-常规程序）.xlsx"
)
FE_MODEL = (
    REPO_ROOT
    / "audit-platform"
    / "frontend"
    / "src"
    / "components"
    / "workpaper"
    / "composables"
    / "d4DisclosureModel.ts"
)
FE_SEGMENT = FE_MODEL.parent / "d4RevenueSegmentColumns.ts"

# 源模板定位（sheet 名逐字取自 wb.sheetnames；括号写法禁统一）
VARIANTS = {
    "listed": {
        "sheet": "附注披露信息（上市公司）",
        "section_title_cell": "A44",
        "section_title": "（4）营业收入、营业成本按分解信息",
        "cat_header_row": 46,
        "leaf_header_row": 47,
        "data_first_row": 48,
        "total_row": 56,
        "template": "note_template_listed.json",
        "section_number": "五、62",
        "table_name": "营业收入、营业成本按分解信息",
        "label_header": "项目",
    },
    "soe": {
        "sheet": "附注披露信息（国企）",
        "section_title_cell": "A37",
        "section_title": "（4）营业收入分解信息",
        "cat_header_row": 39,
        "leaf_header_row": 40,
        "data_first_row": 41,
        "total_row": 49,
        "template": "note_template_soe.json",
        "section_number": "八、64",
        "table_name": "营业收入分解信息",
        "label_header": "合同分类/报告分部",
    },
}

# 行标签序列（逐字，含缩进空格；空串 = 源模板的空可扩行）
EXPECTED_ROW_LABELS = [
    "主营业务",
    "其中：在某一时点确认",
    "      在某一时段确认",
    "",
    "其他业务",
    "其中：在某一时点确认",
    "      在某一时段确认",
    "      租赁收入",
    "合  计",
]

EXPECTED_CATEGORY_LABELS = ["消费品", "汽车", "能源", "其他"]


@pytest.fixture(scope="module")
def wb():
    assert SRC_XLSX.exists(), f"源模板缺失：{SRC_XLSX}"
    return load_workbook(SRC_XLSX, data_only=False)


def _norm(v) -> str:
    return "" if v is None else str(v)


class TestSourceTemplateFacts:
    """源 xlsx 事实（本组失败 = 源模板变了，须重新裁决而不是改代码）。"""

    def test_both_sheets_exist(self, wb):
        for variant, spec in VARIANTS.items():
            assert spec["sheet"] in wb.sheetnames, (
                f"{variant}: sheet {spec['sheet']!r} 不存在；实际 {wb.sheetnames}"
            )

    def test_section_title_literal(self, wb):
        for variant, spec in VARIANTS.items():
            ws = wb[spec["sheet"]]
            got = _norm(ws[spec["section_title_cell"]].value).strip()
            assert got == spec["section_title"], (
                f"{variant}: {spec['section_title_cell']} 实为 {got!r}，"
                f"期望 {spec['section_title']!r}"
            )

    def test_category_header_has_four_merged_pairs(self, wb):
        """类别表头行有 4 个跨 2 列的合并区（B/D/F/H 起）。"""
        for variant, spec in VARIANTS.items():
            ws = wb[spec["sheet"]]
            row = spec["cat_header_row"]
            merged = [
                m
                for m in ws.merged_cells.ranges
                if m.min_row == row and m.max_row == row and m.min_col >= 2
            ]
            assert len(merged) == 4, (
                f"{variant}: 类别表头行 {row} 合并区 {len(merged)} 个（期望 4）："
                f"{[str(m) for m in merged]}"
            )
            for m in merged:
                assert m.max_col - m.min_col == 1, f"{variant}: {m} 未跨 2 列"
            starts = sorted(m.min_col for m in merged)
            assert starts == [2, 4, 6, 8], f"{variant}: 类别起始列 {starts}"

    def test_category_labels_are_the_four_examples(self, wb):
        for variant, spec in VARIANTS.items():
            ws = wb[spec["sheet"]]
            row = spec["cat_header_row"]
            got = [_norm(ws.cell(row, c).value).strip() for c in (2, 4, 6, 8)]
            assert got == EXPECTED_CATEGORY_LABELS, f"{variant}: 类别名 {got}"

    def test_leaf_header_is_revenue_cost_pairs(self, wb):
        for variant, spec in VARIANTS.items():
            ws = wb[spec["sheet"]]
            row = spec["leaf_header_row"]
            got = [_norm(ws.cell(row, c).value).strip() for c in range(2, 10)]
            assert got == ["收入", "成本"] * 4, f"{variant}: 叶子表头 {got}"

    def test_no_horizontal_total_column(self, wb):
        """🔴 关键实证：该表**没有**横向合计列 —— 合计是最后一行。"""
        for variant, spec in VARIANTS.items():
            ws = wb[spec["sheet"]]
            # 类别表头行与叶子表头行在 J 列（第 10 列）及其后都不得有内容
            for row in (spec["cat_header_row"], spec["leaf_header_row"]):
                for c in range(10, min(ws.max_column, 14) + 1):
                    v = _norm(ws.cell(row, c).value).strip()
                    assert v == "", (
                        f"{variant}: 第 {row} 行第 {c} 列有内容 {v!r} ⇒ 可能存在合计列，"
                        "需重新裁决"
                    )
            # 合计**行**确实存在
            total_label = _norm(ws.cell(spec["total_row"], 1).value).strip()
            assert "合" in total_label and "计" in total_label, (
                f"{variant}: 第 {spec['total_row']} 行首列 {total_label!r} 不是合计行"
            )

    def test_row_labels_match_expected_sequence(self, wb):
        for variant, spec in VARIANTS.items():
            ws = wb[spec["sheet"]]
            first = spec["data_first_row"]
            got = [
                _norm(ws.cell(first + i, 1).value)
                for i in range(len(EXPECTED_ROW_LABELS))
            ]
            assert got == EXPECTED_ROW_LABELS, (
                f"{variant}: 行标签序列不符\n实际 {got}\n期望 {EXPECTED_ROW_LABELS}"
            )

    def test_subtotal_formulas_cover_three_rows_including_blank(self, wb):
        """🔴 父行 SUM 覆盖 3 行 —— 含那个空可扩行（这是它必须计入小计的依据）。"""
        for variant, spec in VARIANTS.items():
            ws = wb[spec["sheet"]]
            first = spec["data_first_row"]
            for offset, span_start_off, span_end_off in ((0, 1, 3), (4, 5, 7)):
                cell = ws.cell(first + offset, 2)
                formula = _norm(cell.value)
                assert formula.startswith("=SUM("), (
                    f"{variant}: 第 {first + offset} 行 B 列不是 SUM 公式：{formula!r}"
                )
                m = re.match(r"=SUM\(B(\d+):B(\d+)\)", formula)
                assert m, f"{variant}: SUM 形态异常 {formula!r}"
                lo, hi = int(m.group(1)), int(m.group(2))
                assert lo == first + span_start_off, f"{variant}: SUM 起 {lo}"
                assert hi == first + span_end_off, f"{variant}: SUM 止 {hi}"
                assert hi - lo == 2, f"{variant}: SUM 未覆盖 3 行（{formula}）"

    def test_total_row_is_sum_of_two_subtotals(self, wb):
        for variant, spec in VARIANTS.items():
            ws = wb[spec["sheet"]]
            formula = _norm(ws.cell(spec["total_row"], 2).value)
            first = spec["data_first_row"]
            main_row, other_row = first, first + 4
            assert formula in (
                f"=B{other_row}+B{main_row}",
                f"=B{main_row}+B{other_row}",
            ), f"{variant}: 合计行公式 {formula!r}"


class TestNoteTemplateAlignment:
    """附注模板 JSON 与源 xlsx 对齐（`fix_note_d4_segment_structure.py` 的成果）。"""

    def _table(self, variant: str) -> dict:
        spec = VARIANTS[variant]
        path = REPO_ROOT / "backend" / "data" / spec["template"]
        data = json.loads(path.read_text(encoding="utf-8"))
        secs = [
            s
            for s in data.get("sections", [])
            if (s.get("section_number") or "").strip() == spec["section_number"]
        ]
        assert len(secs) == 1, f"{variant}: 章节命中 {len(secs)} 条"
        tables = [
            t for t in (secs[0].get("tables") or []) if t.get("name") == spec["table_name"]
        ]
        assert len(tables) == 1, f"{variant}: 子表命中 {len(tables)} 条"
        return tables[0]

    @pytest.mark.parametrize("variant", sorted(VARIANTS))
    def test_nine_columns_no_total(self, variant):
        t = self._table(variant)
        cols = t.get("columns") or []
        assert len(cols) == 9, f"{variant}: columns {len(cols)} 个（期望 9）"
        keys = [c.get("key") for c in cols]
        assert "total_revenue" not in keys and "total_cost" not in keys, (
            f"{variant}: 出现横向合计列 ⇒ 比源模板多列：{keys}"
        )

    @pytest.mark.parametrize("variant", sorted(VARIANTS))
    def test_column_keys_are_stable_form(self, variant):
        t = self._table(variant)
        cols = t.get("columns") or []
        assert cols[0]["key"] == "label"
        assert cols[0].get("flat") is True, "标签列须显式 flat（否则被推断出父表头）"
        for c in cols[1:]:
            assert re.fullmatch(r"cat_\d+_(revenue|cost)", c["key"]), (
                f"{variant}: 列 key {c['key']!r} 不是 `cat_N_{{revenue|cost}}` 形态"
            )
            assert c.get("group"), f"{variant}: 列 {c['key']} 缺 group（两级表头）"

    @pytest.mark.parametrize("variant", sorted(VARIANTS))
    def test_column_groups_match_source_categories(self, variant):
        t = self._table(variant)
        groups = []
        for c in (t.get("columns") or [])[1:]:
            if c["group"] not in groups:
                groups.append(c["group"])
        assert groups == EXPECTED_CATEGORY_LABELS, f"{variant}: group 序列 {groups}"

    @pytest.mark.parametrize("variant", sorted(VARIANTS))
    def test_label_header_differs_by_variant(self, variant):
        """两版标签列头不同（上市「项目」/ 国企「合同分类/报告分部」），禁统一。"""
        t = self._table(variant)
        assert t["columns"][0]["label"] == VARIANTS[variant]["label_header"]

    def test_two_variants_label_header_are_different(self):
        a = self._table("listed")["columns"][0]["label"]
        b = self._table("soe")["columns"][0]["label"]
        assert a != b, "两版标签列头被统一了 —— 源模板不同，禁统一"

    @pytest.mark.parametrize("variant", sorted(VARIANTS))
    def test_rows_match_source(self, variant):
        t = self._table(variant)
        got = [r.get("label", "") for r in (t.get("rows") or [])]
        assert got == EXPECTED_ROW_LABELS, f"{variant}: rows {got}"

    @pytest.mark.parametrize("variant", sorted(VARIANTS))
    def test_row_types_mark_subtotal_and_total(self, variant):
        t = self._table(variant)
        rows = t.get("rows") or []
        kinds = [r.get("row_type") for r in rows]
        assert kinds[0] == "subtotal", f"{variant}: 主营业务行 row_type={kinds[0]!r}"
        assert kinds[4] == "subtotal", f"{variant}: 其他业务行 row_type={kinds[4]!r}"
        assert kinds[-1] == "total", f"{variant}: 合计行 row_type={kinds[-1]!r}"

    @pytest.mark.parametrize("variant", sorted(VARIANTS))
    def test_guidance_present(self, variant):
        t = self._table(variant)
        g = t.get("guidance") or ""
        assert len(g) >= 40, f"{variant}: guidance 过短（{len(g)} 字）"
        assert "**" not in g, "guidance 禁 markdown 粗体（TAB 与 Word 都不解析）"


class TestFrontendCrossLock:
    """交叉锁死前端源码：列构造不得再产出横向合计列，行集真源唯一。"""

    def test_model_has_no_total_column(self):
        src = FE_MODEL.read_text(encoding="utf-8")
        # 剥块注释与行注释（说明文字里会写反例）
        code = re.sub(r"/\*[\s\S]*?\*/", "", src)
        code = re.sub(r"^\s*//.*$", "", code, flags=re.M)
        assert "total_revenue" not in code, (
            "d4DisclosureModel.ts 仍产出 total_revenue 列 ⇒ 附注会比源模板多列"
        )
        assert "total_cost" not in code

    def test_model_builds_columns_from_categories(self):
        src = FE_MODEL.read_text(encoding="utf-8")
        assert "buildD4TransposeColumns" in src
        assert "${cat.key}_revenue" in src, "列 key 未由类别 key 派生"
        assert "group: cat.label" in src, "未按类别 label 声明 group（两级表头）"

    def test_row_set_single_source(self):
        src = FE_SEGMENT.read_text(encoding="utf-8")
        assert "D4_SEGMENT_ROWS" in src, "行集真源常量缺失"
        for label in EXPECTED_ROW_LABELS:
            if not label:
                continue
            assert label in src, f"行集真源缺行标签 {label!r}"

    def test_deprecated_constant_marked(self):
        """旧 3 项常量必须标 @deprecated（防新代码继续用它）。"""
        src = FE_MODEL.read_text(encoding="utf-8")
        i = src.find("D4_TRANSPOSE_CHECK_ITEMS")
        assert i > 0
        head = src[max(0, i - 900) : i]
        assert "@deprecated" in head, "D4_TRANSPOSE_CHECK_ITEMS 未标记退役"

    def test_reverse_self_check_scan_is_not_empty(self):
        """反向自检：两个前端文件确实被读到（否则上面的断言全是空转）。"""
        assert len(FE_MODEL.read_text(encoding="utf-8")) > 2000
        assert len(FE_SEGMENT.read_text(encoding="utf-8")) > 2000

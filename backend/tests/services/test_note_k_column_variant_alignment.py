"""K 循环披露表列头分变体、源模板缺陷登记、K11 符号翻转（Requirement 8.2~8.7 / Task 17）。

Task 16 已锁死 8.1（26 个 sheet 名逐字）与列结构自洽。本文件收剩下四件：

- **8.2**：两版用语不同的表**必须分变体声明**，不得共用一份列头。
  判据是「已实测不同的那 12 张表，两版列头必须**保持不同**」—— 反向锁死
  「顺手统一成一份常量」这个改法（那会让一侧的用语整体错掉）。
- **8.5**：`flat`/`group` 在 seed（模板 JSON）与推送（前端载荷 columns）**两侧都表态**；
  不推 columns 的循环不强制（那种情况下附注列头来自模板 seed，本就没有第二侧）。
- **8.6**：源模板自身缺陷**按意图实现并登记**，不照抄。两处逐格实证：
  K3 soe 按性质表标签列列头源值是 `账  龄`（listed 是 `项 目` 才对）；
  K1 两版 `#REF!` 断链（Task 3 已冻结计数，此处只复用不重数）。
- **8.7**：K11 注文「明细表按正数填列、披露表按负数填列」必须在**载荷层**实现。
  实测原状：仅注释与界面文案提到「负数」，**代码里一处翻转都没有** ⇒ 附注里的
  资产减值损失会以正数出现，与利润表口径相反。

spec: .kiro/specs/k-cycle-extraction-formula-and-disclosure-closure/
      Requirements 8.2, 8.4, 8.5, 8.6, 8.7
"""

from __future__ import annotations

import importlib.util
import json
import re
import sys
from pathlib import Path

import pytest

_ROOT = Path(__file__).resolve().parents[3]
_DATA = _ROOT / "backend/data"
_TPL_DIR = _ROOT / "backend/wp_templates/K"
_FE = _ROOT / "audit-platform/frontend/src/components/workpaper/composables"


def _facts():
    spec = importlib.util.spec_from_file_location(
        "_k_facts_17", _ROOT / "backend/tests/test_k_source_template_facts.py"
    )
    assert spec is not None and spec.loader is not None
    mod = importlib.util.module_from_spec(spec)
    sys.modules["_k_facts_17"] = mod
    try:
        spec.loader.exec_module(mod)
    finally:
        sys.modules.pop("_k_facts_17", None)
    return mod


_FACTS = _facts()


@pytest.fixture(scope="module")
def templates() -> dict[str, dict]:
    return {
        v: json.loads((_DATA / f"note_template_{v}.json").read_text(encoding="utf-8"))
        for v in ("listed", "soe")
    }


@pytest.fixture(scope="module")
def k_sections() -> dict[str, dict[str, str]]:
    doc = json.loads((_DATA / "note_workpaper_sync_registry.json").read_text(encoding="utf-8"))
    return {
        str(e["wp_code"]).upper(): {
            "listed": str(e.get("listed") or ""),
            "soe": str(e.get("soe") or ""),
        }
        for e in doc.get("entries") or []
        if str(e.get("wp_code") or "").upper().startswith("K")
    }


def _section(templates: dict, variant: str, number: str) -> dict | None:
    for sec in templates[variant].get("sections") or []:
        if isinstance(sec, dict) and str(sec.get("section_number") or "").strip() == number:
            return sec
    return None


def _labels_by_table(templates: dict, variant: str, number: str) -> dict[str, list[str]]:
    sec = _section(templates, variant, number)
    return {
        str(t.get("name") or ""): [str(c.get("label") or "") for c in (t.get("columns") or [])]
        for t in ((sec or {}).get("tables") or [])
        if isinstance(t, dict)
    }


# ─────────────────────────────────────────────────────────────────────────────
# 8.2：两版列头分变体
# ─────────────────────────────────────────────────────────────────────────────

#: 实测「两版同名表但列头**确实不同**」的清单（2026-08-12 逐表实扫）。
#:
#: 🔴 判据方向是「必须保持不同」：这些表的两版用语真的不一样
#: （`上年年末余额` vs `期初余额`、`项目（或被投资单位）` vs `项目`、
#:  K3 soe 主表标签列是 `类别`、K5/K7/K10 两版列数都不同…），
#: 一旦有人「统一成一份常量」，一侧的用语就整体错掉，而列数相同的那几张
#: **不会有任何报错**。故必须正向锁死差异。
VARIANT_DIVERGENT_TABLES: dict[str, tuple[str, ...]] = {
    "K1": ("其他应收款", "应收利息分类", "应收股利"),
    "K2": ("其他流动资产",),
    "K3": ("其他应付款", "应付利息", "应付股利"),
    "K4": ("其他流动负债",),
    "K5": ("预计负债",),
    "K6": ("持有待售资产减值准备",),
    "K7": ("递延收益",),
    "K10": ("其他收益",),
}

#: 两版列头**确实相同**的表（合法共用，登记以免被误判成漏分变体）。
VARIANT_IDENTICAL_TABLES: dict[str, tuple[str, ...]] = {
    "K1": ("重要逾期利息",),
    "K11": ("资产减值损失",),
    "K12": ("营业外收入",),
    "K13": ("营业外支出",),
}


def test_divergent_tables_stay_divergent(templates, k_sections):
    """8.2：登记为「两版不同」的表，两版列头必须保持不同。"""
    same: list[str] = []
    checked = 0
    for wp, names in sorted(VARIANT_DIVERGENT_TABLES.items()):
        lt = _labels_by_table(templates, "listed", k_sections[wp]["listed"])
        st = _labels_by_table(templates, "soe", k_sections[wp]["soe"])
        for name in names:
            assert name in lt, f"{wp} listed 缺表 {name!r}（表名漂移？）"
            assert name in st, f"{wp} soe 缺表 {name!r}（表名漂移？）"
            checked += 1
            if lt[name] == st[name]:
                same.append(f"{wp}/{name}: 两版列头变成一样了 {lt[name]}")
    assert checked >= 12, f"只核到 {checked} 张表，登记表可疑"
    assert not same, (
        "8.2 违规：这些表两版用语本来不同，被统一成一份了 —— 一侧用语会整体错掉，"
        f"且列数相同时不会有任何报错：{same}"
    )


def test_identical_tables_stay_identical(templates, k_sections):
    """反向自检：登记为「两版相同」的表若变成不同，说明有人只改了一侧。"""
    diff: list[str] = []
    checked = 0
    for wp, names in sorted(VARIANT_IDENTICAL_TABLES.items()):
        lt = _labels_by_table(templates, "listed", k_sections[wp]["listed"])
        st = _labels_by_table(templates, "soe", k_sections[wp]["soe"])
        for name in names:
            if name not in lt or name not in st:
                diff.append(f"{wp}/{name}: 一侧缺表")
                continue
            checked += 1
            if lt[name] != st[name]:
                diff.append(f"{wp}/{name}: L={lt[name]} S={st[name]}")
    assert checked >= 4, f"只核到 {checked} 张表，登记表可疑"
    assert not diff, f"登记为两版相同的表出现分叉（只改了一侧？）：{diff}"


def test_registries_do_not_overlap():
    """两份登记表不得有交集（同一张表不能既「必须不同」又「必须相同」）。"""
    for wp in set(VARIANT_DIVERGENT_TABLES) & set(VARIANT_IDENTICAL_TABLES):
        dup = set(VARIANT_DIVERGENT_TABLES[wp]) & set(VARIANT_IDENTICAL_TABLES[wp])
        assert not dup, f"{wp} 的 {sorted(dup)} 同时出现在两份登记表里"


# ─────────────────────────────────────────────────────────────────────────────
# 8.4 / 8.5：flat 两侧表态
# ─────────────────────────────────────────────────────────────────────────────


def test_seed_side_single_level_tables_declare_flat(templates, k_sections):
    """8.4：模板侧单级表的标签列必须显式标 `flat`。

    不标会让 `_infer_groups_from_headers` 推出**凭空父表头**。
    """
    bad: list[str] = []
    checked = 0
    for wp, nums in sorted(k_sections.items()):
        for variant in ("listed", "soe"):
            num = nums[variant]
            if not num:
                continue
            sec = _section(templates, variant, num)
            for t in ((sec or {}).get("tables") or []):
                cols = t.get("columns") or []
                if not cols or t.get("_column_groups"):
                    continue  # 两级表头由 group 表达
                checked += 1
                if not cols[0].get("flat"):
                    bad.append(f"{wp}/{variant} §{num} / {t.get('name')!r}")
    assert checked >= 40, f"只核到 {checked} 张单级表，判据源可疑"
    assert not bad, f"8.4 违规：单级表标签列未标 flat（会推出凭空父表头）：{bad[:8]}"


#: 前端**确实推 columns** 的循环（其余循环列头由模板 seed 提供，无第二侧可表态）。
FE_COLUMN_PUSHERS: tuple[str, ...] = ("K2", "K3", "K4", "K5", "K6", "K7")


def test_push_side_declares_flat_or_group():
    """8.5：推 columns 的循环，前端也必须表态 `flat`/`group`（不能只在 seed 侧表）。"""
    bad: list[str] = []
    for wp in FE_COLUMN_PUSHERS:
        path = _FE / f"{wp.lower()}NoteSectionMap.ts"
        assert path.exists(), f"前端 map 缺失：{path}"
        src = path.read_text(encoding="utf-8")
        flat = len(re.findall(r"\bflat\s*:", src))
        group = len(re.findall(r"\bgroup\s*:", src))
        if flat + group == 0:
            bad.append(f"{wp}: flat/group 一处都没表态")
    assert not bad, f"8.5 违规（推送侧不表态 ⇒ 附注列头可能被推成凭空两级）：{bad}"


def test_non_pushers_really_do_not_push_columns():
    """反向自检：未列入 `FE_COLUMN_PUSHERS` 的循环确实不推 columns。

    否则上一条会漏掉真正需要表态的循环（清单与代码脱钩 = 判据空转）。
    """
    stray: list[str] = []
    for wp in _FACTS.DISCLOSURE_SHEETS:
        if wp in FE_COLUMN_PUSHERS:
            continue
        path = _FE / f"{wp.lower()}NoteSectionMap.ts"
        if not path.exists():
            continue
        src = path.read_text(encoding="utf-8")
        if re.search(r"\bcolumns\s*:", src):
            stray.append(wp)
    assert not stray, (
        f"这些循环其实推了 columns 却没登记进 FE_COLUMN_PUSHERS：{stray}"
    )


# ─────────────────────────────────────────────────────────────────────────────
# 8.6：源模板缺陷按意图实现并登记
# ─────────────────────────────────────────────────────────────────────────────

#: 源模板缺陷登记（逐格 openpyxl 实证）。
#:
#: 形态：``{key: (wp, variant, coord, 源值, 应实现值, 理由)}``
SOURCE_DEFECTS: dict[str, tuple[str, str, str, str, str, str]] = {
    "k3_soe_nature_table_label_header": (
        "K3",
        "soe",
        "A12",
        "账  龄",
        "项目",
        "「按款项性质列示」表的标签列列头，源模板 soe 侧误写成账龄表的表头；"
        "listed 侧同一张表是「项 目」才对。按 listed 字面实现，源 xlsx 不改。",
    ),
}


@pytest.mark.parametrize("key", sorted(SOURCE_DEFECTS))
def test_source_defect_still_reproduces(key: str):
    """缺陷在源 xlsx 里仍复现 —— stale 检测：源模板修好了要打红提醒移出登记。"""
    from openpyxl import load_workbook

    wp, variant, coord, raw, _intent, _why = SOURCE_DEFECTS[key]
    sheet = _FACTS.DISCLOSURE_SHEETS[wp][variant]
    path = next(iter(sorted(_TPL_DIR.glob(f"{wp} *.xlsx"))))
    wb = load_workbook(path, read_only=False, data_only=True)
    try:
        got = str(wb[sheet][coord].value or "")
    finally:
        wb.close()
    assert got == raw, (
        f"源模板 {wp}/{variant}/{coord} 现值 {got!r} != 登记值 {raw!r}\n"
        "若源模板已修好，请把本条从 SOURCE_DEFECTS 移出（登记不得长期失效）。"
    )


@pytest.mark.parametrize("key", sorted(SOURCE_DEFECTS))
def test_template_implements_intent_not_verbatim(templates, k_sections, key: str):
    """模板侧必须按**意图**实现，不得照抄源模板的错值。"""
    wp, variant, _coord, raw, intent, _why = SOURCE_DEFECTS[key]
    labels = {
        name: cols
        for name, cols in _labels_by_table(templates, variant, k_sections[wp][variant]).items()
    }
    first_cols = {name: (cols[0] if cols else "") for name, cols in labels.items()}
    assert intent in first_cols.values(), (
        f"{wp}/{variant} 没有任何表的标签列列头是意图值 {intent!r}：{first_cols}"
    )
    copied = [name for name, first in first_cols.items() if first == raw]
    assert not copied, f"{wp}/{variant} 照抄了源模板错值 {raw!r}：{copied}"


def test_every_defect_has_substantial_reason():
    for key, (_wp, _v, _c, _raw, _intent, why) in SOURCE_DEFECTS.items():
        assert len(why) >= 25, f"{key} 登记理由过短，读者看不出为什么不照抄：{why!r}"


def test_k1_ref_errors_are_registered_not_recopied():
    """K1 `#REF!` 断链沿用 Task 3 的冻结计数（此处只复用，不重新数）。"""
    counts = _FACTS.K1_REF_ERROR_COUNTS
    assert sum(counts.values()) > 0, "K1 `#REF!` 计数为 0 ⇒ 登记失效或判据源变了"
    assert set(counts) == {"listed", "soe"}


# ─────────────────────────────────────────────────────────────────────────────
# 8.7：K11 披露表按负数填列
# ─────────────────────────────────────────────────────────────────────────────

_K11_MAP = _FE / "k11NoteSectionMap.ts"


def test_k11_sign_flip_exists_in_payload_layer():
    """8.7：翻转必须在**载荷构造**里，不能只写在注释/界面文案里。

    实测原状就是「只有文案说负数、代码一处翻转都没有」—— 那种状态下附注里的
    资产减值损失以正数出现，与利润表口径相反，而任何测试都不会红。
    """
    src = _K11_MAP.read_text(encoding="utf-8")
    assert "export function k11DisclosureAmount" in src, (
        "8.7 未实现：k11NoteSectionMap.ts 里没有符号翻转函数"
    )
    # 翻转必须真的作用在载荷行上
    assert re.search(r"currentAmount:\s*k11DisclosureAmount\(", src), (
        "8.7 未接线：本期发生额没有过翻转函数"
    )
    assert re.search(r"priorAmount:\s*k11DisclosureAmount\(", src), (
        "8.7 未接线：上期发生额没有过翻转函数"
    )


def test_k11_flip_is_single_point():
    """翻转只许一处 —— 组件里再翻一次会翻回正数且**不会报错**。"""
    comp_dir = _ROOT / "audit-platform/frontend/src/components/workpaper"
    offenders: list[str] = []
    for path in sorted(comp_dir.rglob("K11*.vue")):
        text = path.read_text(encoding="utf-8")
        for i, line in enumerate(text.splitlines(), 1):
            stripped = line.strip()
            if stripped.startswith(("*", "//", "<!--")):
                continue
            if re.search(r"(?:\*\s*-1|-1\s*\*|k11DisclosureAmount\()", stripped):
                offenders.append(f"{path.name}:{i}: {stripped[:90]}")
    assert not offenders, f"组件里也在翻符号 ⇒ 双重翻转（翻回正数）：{offenders}"


def test_k11_source_note_really_says_negative():
    """反向自检：源 xlsx 两版确实写着「披露表按照负数填列」。

    没有这条，上面两条就是「按注释办事」——而注释可能本身就是抄错的。
    """
    from openpyxl import load_workbook

    path = next(iter(sorted(_TPL_DIR.glob("K11 *.xlsx"))))
    wb = load_workbook(path, read_only=True, data_only=True)
    try:
        hits: list[str] = []
        for variant, sheet in _FACTS.DISCLOSURE_SHEETS["K11"].items():
            ws = wb[sheet]
            for row in ws.iter_rows():
                for c in row:
                    if isinstance(c.value, str) and "负数填列" in c.value:
                        hits.append(f"{variant}/{c.coordinate}")
    finally:
        wb.close()
    assert len(hits) >= 2, (
        f"源 xlsx 里找不到「负数填列」注文（两版各应至少 1 处），实测 {hits}\n"
        "若源模板改了口径，8.7 的实现要跟着改而不是留着。"
    )

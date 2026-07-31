"""附注「递延所得税资产和递延所得税负债」两节**行集**对齐源模板守卫。

spec: `.kiro/specs/n1-four-table-extraction-and-disclosure-alignment/`
Properties: 9（行集逐字等于源模板）/ 10（无占位假行 + 科目与报表行白名单）/ 11（幂等）

为什么单独一个文件
------------------
既有 `test_note_deferred_tax_structure.py` 只校验**列结构 / 两级表头 / guidance / 小计标记**，
对 `rows[].label` 完全没有断言 → 上市表 1 曾长期携带源模板不存在的「开办费」、
负债段 5 行里 4 行不符、国企残留 5 处 `……` 占位假行而全部测试通过。
本文件用 **openpyxl 直读源 xlsx** 做交叉比对，把行集钉死。

🔴 直读源模板而非引用脚本常量，否则守卫退化成「脚本自己和自己比」（空转）。
"""

from __future__ import annotations

import json
from pathlib import Path

import pytest
from openpyxl import load_workbook

BACKEND = Path(__file__).resolve().parents[2]
DATA_DIR = BACKEND / "data"
SRC_XLSX = BACKEND / "wp_templates" / "N" / "N1 递延所得税资产.xlsx"

SHEET_LISTED = "附注披露信息（上市公司）"
SHEET_SOE = "附注披露信息（国企）"

SECTIONS = {"listed": "五、30", "soe": "八、31"}
PATHS = {
    "listed": DATA_DIR / "note_template_listed.json",
    "soe": DATA_DIR / "note_template_soe.json",
}

T_UNOFFSET = "未经抵销的递延所得税资产和递延所得税负债"
T_NET_OFFSET = "以抵销后净额列示的递延所得税资产或负债"
T_OFFSET_DETAIL = "递延所得税资产和递延所得税负债互抵明细"
T_LOSS_EXPIRY = "未确认递延所得税资产的可抵扣亏损将于以下年度到期"

# 源模板单元格区段（逐格实测：listed A1:K54 / soe A1:IV74）
SRC_RANGES = {
    # 上市：资产段 R13:R19（7 项）、负债段 R23:R27（5 项）
    "listed_asset": (SHEET_LISTED, 13, 19),
    "listed_liability": (SHEET_LISTED, 23, 27),
    # 国企：资产段 R13:R19 是**公式**引用上市同列（`='附注披露信息（上市公司）'!A13`）
    #       → 语义上两版资产段同集，故不逐格读，改为断言等于上市资产段。
    "soe_liability": (SHEET_SOE, 23, 27),
    # 亏损到期年度：上市 R46:R51 / 国企 R66:R71，各 6 行
    "listed_loss": (SHEET_LISTED, 46, 51),
    "soe_loss": (SHEET_SOE, 66, 71),
    # 国企（2）B 互抵明细：R56:R58 全空（纯动态行区域）
    "soe_offset_detail": (SHEET_SOE, 56, 58),
}


# ─── fixtures ────────────────────────────────────────────────────────────────


@pytest.fixture(scope="module")
def wb():
    assert SRC_XLSX.exists(), f"源模板缺失：{SRC_XLSX}（运行时权威目录）"
    return load_workbook(SRC_XLSX, data_only=False)


@pytest.fixture(scope="module")
def sections() -> dict[str, dict]:
    out: dict[str, dict] = {}
    for variant, path in PATHS.items():
        data = json.loads(path.read_text(encoding="utf-8"))
        hits = [s for s in data["sections"] if s.get("section_number") == SECTIONS[variant]]
        assert len(hits) == 1, f"{variant} {SECTIONS[variant]} 应恰好 1 个 section"
        out[variant] = hits[0]
    return out


def _src_labels(wb, key: str) -> list[str]:
    sheet, r0, r1 = SRC_RANGES[key]
    ws = wb[sheet]
    out: list[str] = []
    for r in range(r0, r1 + 1):
        v = ws.cell(row=r, column=1).value
        out.append("" if v is None else str(v).strip())
    return out


def _table(section: dict, name: str) -> dict:
    hits = [t for t in (section.get("tables") or []) if t.get("name") == name]
    assert len(hits) == 1, f"表 {name!r} 应恰好 1 张，实际 {len(hits)}"
    return hits[0]


def _labels(table: dict) -> list[str]:
    return [str(r.get("label", "")).strip() for r in (table.get("rows") or [])]


def _data_labels(table: dict) -> list[str]:
    """只取明细数据行（排除小计 / 合计 / 分组标题）。"""
    out: list[str] = []
    for r in table.get("rows") or []:
        if r.get("is_total"):
            continue
        label = str(r.get("label", "")).strip()
        # 分组标题行携带 report_row_code（承载 BS 报表行映射）
        if r.get("report_row_code"):
            continue
        out.append(label)
    return out


# ─── 反向自检：先证明源模板读取本身非空转 ─────────────────────────────────────


def test_self_check_source_ranges_are_non_empty(wb):
    """若源模板路径/区段写错，下面所有比对都会退化成空对空 → 先钉死读取有效性。"""
    assert _src_labels(wb, "listed_asset") == [
        "资产减值准备",
        "可抵扣亏损",
        "内部交易未实现利润",
        "公允价值变动",
        "租赁负债",
        "购入摊销年限小于税法规定的资产",
        "其他",
    ]
    assert _src_labels(wb, "listed_liability")[0] == "购入摊销年限大于税法规定的资产"
    assert _src_labels(wb, "soe_liability")[3] == "租赁形成", "国企 A26 应为「租赁形成」"
    assert _src_labels(wb, "listed_liability")[3] == "使用权资产", "上市 A26 应为「使用权资产」"
    # 两版负债段第 4 项必须不同（若相同，本 spec 的分变体逻辑就失去意义）
    assert _src_labels(wb, "listed_liability")[3] != _src_labels(wb, "soe_liability")[3]


def test_self_check_soe_asset_segment_mirrors_listed(wb):
    """国企资产段是公式引用上市同列 → 证明「两版资产段同集」的依据成立。"""
    ws = wb[SHEET_SOE]
    for r in range(13, 20):
        f = str(ws.cell(row=r, column=1).value or "")
        assert "附注披露信息（上市公司）" in f, f"国企 A{r} 应为引用上市的公式，实际 {f!r}"


# ─── Property 9: 行集逐字等于源模板 ──────────────────────────────────────────


def test_listed_unoffset_rows_match_source(wb, sections):
    t = _table(sections["listed"], T_UNOFFSET)
    expected = _src_labels(wb, "listed_asset") + _src_labels(wb, "listed_liability")
    assert _data_labels(t) == expected


def test_soe_unoffset_rows_match_source(wb, sections):
    t = _table(sections["soe"], T_UNOFFSET)
    # 国企资产段 = 上市资产段（公式引用）；负债段取国企自身字面
    expected = _src_labels(wb, "listed_asset") + _src_labels(wb, "soe_liability")
    assert _data_labels(t) == expected


def test_soe_net_offset_rows_mirror_unoffset(wb, sections):
    """源模板国企 A36=`=A13` … A50=`=A27` → 表 2 行集镜像表 1。"""
    t1 = _table(sections["soe"], T_UNOFFSET)
    t2 = _table(sections["soe"], T_NET_OFFSET)
    assert _labels(t2) == _labels(t1)


def test_soe_net_offset_mirror_is_formula_in_source(wb):
    """反向自检：镜像关系的依据是源模板里的公式引用，不是开发假设。"""
    ws = wb[SHEET_SOE]
    for r in (36, 37, 38, 39, 40, 41, 42):
        assert str(ws.cell(row=r, column=1).value or "").startswith("=A"), f"A{r} 应为 =A1x"
    for r in (46, 47, 48, 49, 50):
        assert str(ws.cell(row=r, column=1).value or "").startswith("=A"), f"A{r} 应为 =A2x"


@pytest.mark.parametrize(
    "variant,src_key",
    [("listed", "listed_loss"), ("soe", "soe_loss")],
)
def test_loss_expiry_has_six_year_rows(wb, sections, variant, src_key):
    """源模板 6 个年度行：期末列覆盖 Y+1~Y+5、上年年末列覆盖 Y~Y+4，并集 6 年。"""
    src_years = [x for x in _src_labels(wb, src_key) if x]
    assert len(src_years) == 6, f"源模板 {src_key} 应为 6 个年度行，实际 {src_years}"
    t = _table(sections[variant], T_LOSS_EXPIRY)
    year_rows = [x for x in _labels(t) if x.endswith("年")]
    assert len(year_rows) == 6, f"{variant} 亏损到期表应 6 个年度行，实际 {year_rows}"
    # 年度连续
    nums = [int(x[:-1]) for x in year_rows]
    assert nums == list(range(nums[0], nums[0] + 6)), f"年度应连续，实际 {nums}"
    # 合计行存在且唯一
    assert sum(1 for r in t["rows"] if r.get("is_total")) == 1


def test_soe_offset_detail_is_pure_dynamic_area(wb, sections):
    """源模板（2）B R56:R58 全空 → 模板不得 seed 任何行（否则附注多一行空披露数据）。"""
    assert _src_labels(wb, "soe_offset_detail") == ["", "", ""]
    t = _table(sections["soe"], T_OFFSET_DETAIL)
    assert t.get("rows") == [], f"互抵明细应为空行骨架，实际 {_labels(t)}"
    # 语义必须落在 guidance（不能因删行而丢失「动态行区域」的编制指引）
    assert "动态" in (t.get("guidance") or "")


# ─── Property 10: 无占位假行 + 白名单 ────────────────────────────────────────


@pytest.mark.parametrize("variant", ["listed", "soe"])
def test_no_ellipsis_placeholder_rows(sections, variant):
    """`……` 是 md 重建把源模板「可无限量添加行」当数据行落下的假行 → 必须为 0。"""
    bad: list[str] = []
    for t in sections[variant].get("tables") or []:
        for label in _labels(t):
            if "…" in label or label in {"...", "。。。"}:
                bad.append(f"{t.get('name')}::{label}")
    assert bad == [], f"{variant} 残留占位假数据行：{bad}"


@pytest.mark.parametrize("variant", ["listed", "soe"])
def test_account_codes_whitelist(sections, variant):
    """附注行只许挂递延所得税科目。

    曾错挂 `2601`（租赁负债）/ `1641~1643`（使用权资产）—— 那是**被计量的底层科目**，
    而本表金额是暂时性差异与递延所得税 → 取数会把底层余额拉进递延税列。
    """
    allowed = {"1811", "2901"}
    bad: list[str] = []
    for t in sections[variant].get("tables") or []:
        for r in t.get("rows") or []:
            for c in r.get("account_codes") or []:
                if str(c) not in allowed:
                    bad.append(f"{t.get('name')}::{r.get('label')}::{c}")
    assert bad == [], f"{variant} 行携带非递延所得税科目：{bad}"


@pytest.mark.parametrize("variant", ["listed", "soe"])
def test_report_row_code_whitelist(sections, variant):
    """`BS-036` 递延所得税资产 / `BS-067` 递延所得税负债（report_config DB 实证）。

    曾错写 `BS-018`（上市=流动资产合计 / 国企=存货）与 `BS-042` / `BS-019`。
    """
    allowed = {"BS-036", "BS-067"}
    bad: list[str] = []
    for t in sections[variant].get("tables") or []:
        for r in t.get("rows") or []:
            rc = r.get("report_row_code")
            if rc and rc not in allowed:
                bad.append(f"{t.get('name')}::{r.get('label')}::{rc}")
    assert bad == [], f"{variant} 行携带错误报表行编码：{bad}"


@pytest.mark.parametrize("variant", ["listed", "soe"])
def test_group_header_rows_carry_correct_row_code(sections, variant):
    """资产段标题挂 BS-036、负债段标题挂 BS-067（各恰好一处）。"""
    t = _table(sections[variant], T_UNOFFSET)
    codes = [r.get("report_row_code") for r in t["rows"] if r.get("report_row_code")]
    assert codes == ["BS-036", "BS-067"], f"{variant} 表 1 分组标题报表行应依次为资产/负债，实际 {codes}"


def test_listed_text_sections_have_no_truncated_paragraph(sections):
    """md 抽取遗留的证监会指引段在句中被截断 + 带 `**` 残迹 → 必须已移除。"""
    texts = sections["listed"].get("text_sections") or []
    for t in texts:
        assert "监管规则适用指引" not in t, f"截断段落仍在：{t[:60]}"
        assert "**" not in t, f"text_sections 残留 markdown 加粗标记：{t[:60]}"


# ─── Property 11: 幂等 ───────────────────────────────────────────────────────


def test_fix_script_reports_no_pending_work():
    """`--check` 必须报 0 欠账（本 spec 已 apply）。"""
    import importlib

    mod = importlib.import_module("scripts.fix.fix_note_deferred_tax_structure")
    assert mod.apply(check_only=True) is False, "模板与脚本目标不一致，请重跑 fix 脚本"


def test_build_functions_are_pure():
    """两次 build 结果相等且互不共享可变对象（防行 dict 被跨表复用后就地改写）。"""
    import importlib

    mod = importlib.import_module("scripts.fix.fix_note_deferred_tax_structure")
    for build in (mod.build_listed_tables, mod.build_soe_tables):
        a, b = build(), build()
        assert a == b
        a[0]["rows"][0]["label"] = "__mutated__"
        assert build()[0]["rows"][0]["label"] != "__mutated__"


def test_row_code_constants_are_db_verified():
    import importlib

    mod = importlib.import_module("scripts.fix.fix_note_deferred_tax_structure")
    assert mod.ASSET_ROW_CODE == "BS-036"
    assert mod.LIABILITY_ROW_CODE == "BS-067"
    assert mod.ALLOWED_ACCOUNT_CODES == {"1811", "2901"}

"""批量导出可用性：零 Tab 拦截、模板模式不带项目数据、密码字符集、循环清单下发。

spec: environment-hygiene-deps-and-scratch-schemas（Requirement 8，design §八）
缘起：任务 7.3 真栈复测暴露四个问题（都是「操作看着成功、结果不是用户以为的那样」）：
  ① 导出 0 张底稿时静默成功，界面照样提示「模板 ZIP 已导出」
  ② 「导出空白模板」声称不含项目数据，包内却固定附带报表 / 未审报表 / 附注 / 试算表
  ③ 服务端按 UTF-8 加密，按 GBK 处理密码的解压软件判密码错误
  ④ E、J 在 catalog 里 0 条登记，对话框却可勾选，导出时既不进 files 也不进 skipped

🔴 本文件的判据纪律（design §8.2 明写）：不能只断言「包里没有这些文件」。单测里 `db`
是 `AsyncMock`，四个导出器本就会抛错被 fail-soft 吞掉 ⇒ 附加文件天然不存在，
**删掉 `if mode == "data"` 照样绿**。故把导出器替换成必定成功的替身，断言模板模式
**未调用**、数据模式调用（正向对照）。
"""
from __future__ import annotations

import io
import json
import os
import re
import sys
import uuid
import zipfile
from contextlib import contextmanager
from pathlib import Path
from types import SimpleNamespace
from unittest.mock import AsyncMock, MagicMock, patch

import pytest

os.environ.setdefault("JWT_SECRET_KEY", "test-secret-key-for-unit-tests")

from app.services.bulk_tab import bulk_export_service  # noqa: E402
from app.services.bulk_tab.bulk_export_service import (  # noqa: E402
    PASSWORD_MAX_LENGTH,
    SKIP_REASON_ADVICE,
    SKIP_REASON_LABELS,
    export,
)
from app.services.bulk_tab.exceptions import (  # noqa: E402
    BulkExportNothingToExportError,
    BulkExportPasswordInvalidError,
)
from app.services.bulk_tab.manifest_builder import (  # noqa: E402
    BulkManifest,
    ManifestFileEntry,
    ManifestSkippedEntry,
)

_PATCH_BUILD_MANIFEST = "app.services.bulk_tab.bulk_export_service.build_manifest"
_PATCH_EXPORT_TAB = "app.services.bulk_tab.bulk_export_service.export_tab"

_FRONTEND = Path(__file__).resolve().parents[2].parent / "audit-platform" / "frontend"
_DIALOG = _FRONTEND / "src" / "components" / "workpaper" / "bulk-tab" / "WpBulkDialog.vue"


# ═══ 夹具 ═════════════════════════════════════════════════════════════════════


def _xlsx(text: str = "值") -> bytes:
    import openpyxl

    wb = openpyxl.Workbook()
    wb.active.append(["列1", "列2"])
    wb.active.append([text, "值"])
    buf = io.BytesIO()
    wb.save(buf)
    return buf.getvalue()


def _entry(sheet_code: str = "D2-1") -> ManifestFileEntry:
    return ManifestFileEntry(
        addr_id=f"D2/{sheet_code}", wp_code="D2", parent_wp_code="D2", sheet_code=sheet_code,
        sheet_name=f"明细表{sheet_code}", origin="standard", api_prefix="d2",
        item_id=f"{sheet_code}-vc-rows", storage_field="remark", wp_id=str(uuid.uuid4()),
        import_order=1, depends_on_sheets=[], zip_path=f"D/D2/{sheet_code}_明细表.xlsx", sha256="",
    )


def _skipped(sheet_code: str, reason: str) -> ManifestSkippedEntry:
    return ManifestSkippedEntry(
        addr_id=f"D2/{sheet_code}", sheet_code=sheet_code, sheet_name=f"表{sheet_code}",
        parent_wp_code="D2", skip_reason=reason,
    )


def _manifest(
    *,
    files: list[ManifestFileEntry] | None = None,
    skipped: list[ManifestSkippedEntry] | None = None,
    unsupported: list[str] | None = None,
    mode: str = "template",
    cycles: list[str] | None = None,
) -> BulkManifest:
    return BulkManifest(
        project_id=str(uuid.uuid4()), audit_year=2025, exported_at="2026-09-30T10:00:00Z",
        exported_by="tester", mode=mode, cycles=cycles or ["D"],
        files=list(files or []), skipped=list(skipped or []),
        unsupported_cycles=list(unsupported or []),
    )


def _quiet_db():
    """``db.get`` 返回 None ⇒ 跳过增量清单持久化（与本文件判据无关）。"""
    db = AsyncMock()
    db.get = AsyncMock(return_value=None)
    return db


@contextmanager
def _world(manifest: BulkManifest, *, tab_bytes: bytes | None = None, fail_tabs: bool = False):
    """替换 manifest 构建与逐 Tab 导出；yield build_manifest 替身供「是否开始工作」断言。"""
    payload = tab_bytes if tab_bytes is not None else _xlsx()

    async def _tab(**kwargs):
        if fail_tabs:
            raise RuntimeError("这一张失败")
        return payload

    build = AsyncMock(side_effect=lambda *a, **k: manifest)
    with patch(_PATCH_BUILD_MANIFEST, build), patch(_PATCH_EXPORT_TAB, AsyncMock(side_effect=_tab)):
        yield build


class _SpyReportExporter:
    """必定成功的报表导出器替身 —— 记录调用次数。

    🔴 为什么要替身：真实导出器在 `AsyncMock` db 上必然抛错并被 fail-soft 吞掉，
    于是「模板包里没有 `_报表/`」这件事在删掉 `if mode == "data"` 之后**照样成立** ⇒
    只看包内容的判据是恒绿的。替身让「有没有去调它」变成可观测量。
    """

    calls: list[str] = []

    def __init__(self, db):  # noqa: ANN001
        self._db = db

    async def export(self, *, project_id, year, flatten_formulas=True):  # noqa: ANN001, ANN201
        # 刻意不声明 `mode` 参数：生产代码用 inspect.signature 探测它来决定是否再导未审数
        type(self).calls.append("report")
        return io.BytesIO(_xlsx("报表"))


class _SpyNoteExporter:
    calls: list[str] = []

    def __init__(self, db):  # noqa: ANN001
        self._db = db

    async def export(self, *, project_id, year, template_type):  # noqa: ANN001, ANN201
        type(self).calls.append("note")
        return io.BytesIO(b"docx-bytes")


class _SpyTrialBalance:
    calls: list[str] = []

    def __init__(self, db):  # noqa: ANN001
        self._db = db

    async def get_trial_balance(self, project_id, year):  # noqa: ANN001, ANN201
        type(self).calls.append("tb")
        return [SimpleNamespace(
            standard_account_code="1001", account_name="库存现金", opening_balance=1,
            unadjusted_amount=2, aje_adjustment=0, rje_adjustment=0, audited_amount=2,
        )]


@contextmanager
def _spy_exporters():
    """三个附加文件导出器都换成必定成功的替身，并清空调用记录。"""
    for cls in (_SpyReportExporter, _SpyNoteExporter, _SpyTrialBalance):
        cls.calls = []
    with patch("app.services.report_excel_exporter.ReportExcelExporter", _SpyReportExporter), \
            patch("app.services.note_word_exporter.NoteWordExporter", _SpyNoteExporter), \
            patch("app.services.trial_balance_service.TrialBalanceService", _SpyTrialBalance):
        yield


def _all_spy_calls() -> list[str]:
    return [*_SpyReportExporter.calls, *_SpyNoteExporter.calls, *_SpyTrialBalance.calls]


# ═══ 8.1 零 Tab 即拒绝导出 ═══════════════════════════════════════════════════


@pytest.mark.asyncio
@pytest.mark.parametrize(
    ("reason", "must_contain"),
    [
        ("resolve_instance_miss", "生成底稿"),
        ("no_data", "仅导出有数据的 Tab"),
        ("unchanged", "增量导出"),
        ("export_failed", "导出时出错"),
        ("no_adapter", "暂未接入"),
        ("unknown", "未知原因"),
    ],
)
async def test_zero_tab_reason_is_actionable_chinese(reason, must_contain):
    """每个跳过类别都给一句中文原因，且对可纠正的类别点明下一步怎么做（Req 8.1）。"""
    manifest = _manifest(files=[], skipped=[_skipped("D2-X", reason)])
    with _world(manifest):
        with pytest.raises(BulkExportNothingToExportError) as info:
            await export(db=_quiet_db(), project_id=uuid.uuid4(), cycles=["D"], mode="template")

    message = str(info.value)
    assert info.value.status_code == 422
    assert must_contain in message, message
    # 英文代码不出现在用户可见文案里
    assert reason not in message
    assert info.value.counts.get(reason) == 1


@pytest.mark.asyncio
async def test_zero_tab_does_not_leak_existence_when_all_filtered_by_visibility():
    """全部条目被可见集过滤掉时，只给通用说明，不透露「有多少张被权限挡住」（Req 8.1）。

    可见集过滤掉的条目按 Task 10 设计**不进** `skipped` —— 所以这里 `skipped` 为空，
    错误文案只剩通用句。这条同时锁住「不能为了文案好看而把被过滤的条目也记进 skipped」。
    """
    manifest = _manifest(files=[_entry("D2-1"), _entry("D2-2")])

    async def _deny(wp_id, sheet_code):  # noqa: ANN001, ANN202
        return False

    with _world(manifest):
        with pytest.raises(BulkExportNothingToExportError) as info:
            await export(
                db=_quiet_db(), project_id=uuid.uuid4(), cycles=["D"], mode="template",
                visible_filter=_deny,
            )

    message = str(info.value)
    assert bulk_export_service.NOTHING_VISIBLE_MESSAGE in message
    # 数量、底稿编号、权限字样都不得出现
    assert "2" not in message and "D2-1" not in message
    assert "权限" not in message and "可见" not in message
    assert info.value.counts == {}


@pytest.mark.asyncio
async def test_zero_tab_rejects_before_extra_files_and_incremental_ledger():
    """422 抛在写附加文件与增量清单之前：不做无用功、不留痕迹（Req 8.1）。"""
    db = AsyncMock()
    project = SimpleNamespace(wizard_state={})
    db.get = AsyncMock(return_value=project)

    manifest = _manifest(files=[_entry("D2-1")], mode="data", skipped=[])
    with _world(manifest, fail_tabs=True), _spy_exporters():
        with pytest.raises(BulkExportNothingToExportError):
            await export(db=db, project_id=uuid.uuid4(), cycles=["D"], mode="data")

    # 三个附加文件导出器一个都没被调用
    assert _all_spy_calls() == []
    # 增量清单没被写（wizard_state 里不该出现本次导出的痕迹）
    assert "_last_bulk_export_manifest" not in project.wizard_state
    db.flush.assert_not_awaited()


@pytest.mark.asyncio
async def test_unsupported_cycle_reason_names_the_cycle():
    """整个循环未接入时，原因里点明循环代码与中文名（Req 8.5）。"""
    from app.services.bulk_tab.scenario_registry import UNSUPPORTED_CYCLE_REASON

    manifest = _manifest(files=[], unsupported=["E", "J"], cycles=["E", "J"])
    with _world(manifest):
        with pytest.raises(BulkExportNothingToExportError) as info:
            await export(db=_quiet_db(), project_id=uuid.uuid4(), cycles=["E", "J"], mode="template")

    message = str(info.value)
    assert "E 货币资金" in message and "J 薪酬" in message
    assert UNSUPPORTED_CYCLE_REASON in message
    assert info.value.counts.get("unsupported_cycle") == 2


# ═══ 8.2 模板模式不附带项目数据 ═══════════════════════════════════════════════


@pytest.mark.asyncio
async def test_template_mode_never_calls_project_data_exporters():
    """模板模式**不调用**报表 / 附注 / 试算表导出器（Req 8.3 的反向判据）。"""
    manifest = _manifest(files=[_entry("D2-1")], mode="template")
    with _world(manifest), _spy_exporters():
        result = await export(db=_quiet_db(), project_id=uuid.uuid4(), cycles=["D"], mode="template")

    assert _all_spy_calls() == [], "模板模式不该触碰项目数据导出器"
    with zipfile.ZipFile(result) as zf:
        names = zf.namelist()
    assert not [n for n in names if n.startswith(("_报表/", "_附注/", "_试算表/"))]
    # 底稿目录与 manifest 两种模式都保留
    assert "_底稿目录.xlsx" in names and "manifest.json" in names


@pytest.mark.asyncio
async def test_data_mode_does_call_them_and_readme_lists_what_was_written():
    """正向对照：数据模式**确实**调用三个导出器，且 README 按实际写入逐行列出（Req 8.3）。

    没有这条正向对照，上一个测试在「把 `_write_project_data_files` 整个删掉」这种
    变异下也会绿 —— 那不是修复而是砍功能。
    """
    manifest = _manifest(files=[_entry("D2-1")], mode="data")
    with _world(manifest), _spy_exporters():
        result = await export(db=_quiet_db(), project_id=uuid.uuid4(), cycles=["D"], mode="data")

    assert sorted(_all_spy_calls()) == ["note", "report", "tb"]
    with zipfile.ZipFile(result) as zf:
        names = zf.namelist()
        readme = zf.read("README.txt").decode("utf-8")
    for path in ("_报表/财务报表.xlsx", "_附注/财务报表附注.docx", "_试算表/试算平衡表.xlsx"):
        assert path in names
        assert path in readme, f"README 未列出实际写入的 {path}"
    # 替身的 export 没有 `mode` 参数 ⇒ 未审数那份不会写，README 也不该提它
    assert "_报表/财务报表_未审数.xlsx" not in names
    assert "_报表/财务报表_未审数.xlsx" not in readme


# ═══ 8.3 密码限定可打印 ASCII ════════════════════════════════════════════════


@pytest.mark.asyncio
@pytest.mark.parametrize(
    ("password", "why"),
    [
        ("审计-Pa55!", "中文"),
        ("Ｐａ５５", "全角字母数字"),
        ("pa ss55", "中间空格"),
        (" pass55", "前导空格（输入框里看不见）"),
        ("pass55 ", "尾随空格"),
        ("pass\t55", "制表符"),
        ("pass55\n", "换行"),
        ("密码", "纯中文"),
        ("pass😀", "emoji"),
        ("é" + "pass55", "带音标的拉丁字母"),
        ("x" * (PASSWORD_MAX_LENGTH + 1), "超长"),
    ],
)
async def test_invalid_password_is_422_before_building_manifest(password, why):
    """非法密码在建 manifest **之前**就 422（Req 8.4）。

    「之前」这一点靠 `build.assert_not_awaited()` 而不是靠读代码：校验一旦被挪到
    manifest 之后，导出已经开始读库、增量清单可能已 flush，而用户拿到的仍是同一句错误
    —— 从响应上看不出区别。
    """
    manifest = _manifest(files=[_entry("D2-1")])
    with _world(manifest) as build:
        with pytest.raises(BulkExportPasswordInvalidError) as info:
            await export(
                db=_quiet_db(), project_id=uuid.uuid4(), cycles=["D"],
                mode="template", password=password,
            )

    assert info.value.status_code == 422, why
    assert "英文字母、数字和英文符号" in str(info.value)
    build.assert_not_awaited()


@pytest.mark.asyncio
@pytest.mark.parametrize(
    "password",
    [
        "a",                              # 下边界：1 位
        "x" * PASSWORD_MAX_LENGTH,        # 上边界：恰 128 位
        "Audit-Pa55!",                    # 常规
        "!\"#$%&'()*+,-./:;<=>?@[\\]^_`{|}~",  # 全部 ASCII 符号
        "0123456789",
    ],
)
async def test_legal_password_passes_validation(password):
    """边界合法值不被误拒 —— 否则「修好了安全」会变成「谁都导不出来」。"""
    manifest = _manifest(files=[_entry("D2-1")])
    with _world(manifest), patch.object(bulk_export_service, "_require_zip_encryption"), \
            patch.object(bulk_export_service, "_encrypt_zip", side_effect=lambda z, p: z):
        result = await export(
            db=_quiet_db(), project_id=uuid.uuid4(), cycles=["D"],
            mode="template", password=password,
        )
    assert isinstance(result, io.BytesIO)


@pytest.mark.asyncio
async def test_password_validated_before_encryption_availability():
    """密码违规（422，用户可改）先于「服务器缺 pyzipper」（503，用户改不了）报出。

    顺序反了的后果：用户在密码里打了中文，却被告知「服务器没装组件」——
    去找管理员装依赖，装完还是同一个错。
    """
    saved = sys.modules.get("pyzipper")
    sys.modules["pyzipper"] = None  # type: ignore[assignment]
    try:
        manifest = _manifest(files=[_entry("D2-1")])
        with _world(manifest):
            with pytest.raises(BulkExportPasswordInvalidError):
                await export(
                    db=_quiet_db(), project_id=uuid.uuid4(), cycles=["D"],
                    mode="template", password="审计-Pa55!",
                )
    finally:
        if saved is None:
            sys.modules.pop("pyzipper", None)
        else:
            sys.modules["pyzipper"] = saved


def test_frontend_password_rule_matches_backend_exactly():
    """前后端密码规则逐值一致（Req 8.4 末句「两端规则由守卫交叉锁死」）。

    两份规则各写一遍是必要的（后端权威、前端是体验预检，让用户点按钮前就知道问题），
    但一旦漂移就会出现「前端放行、后端 422」或「前端拦住、后端本来能过」两种都让人
    不信任的状态。本判据读前端源码取常量，与后端常量比对。
    """
    assert _DIALOG.exists(), f"前端对话框不存在：{_DIALOG}"
    src = _DIALOG.read_bytes().decode("utf-8")

    max_len = re.search(r"const PASSWORD_MAX_LENGTH = (\d+)", src)
    assert max_len, "前端未声明 PASSWORD_MAX_LENGTH"
    assert int(max_len.group(1)) == PASSWORD_MAX_LENGTH

    forbidden = re.search(r"const PASSWORD_FORBIDDEN_RE = /(.+?)/", src)
    assert forbidden, "前端未声明 PASSWORD_FORBIDDEN_RE"
    # 两端都用**否定字符类**（无锚点）：Python 的 `$` 还匹配末尾换行之前，
    # 用 `^...+$` 会让两端语义不一致，且放行 `"pass55\n"` 这种必然打不开的密码
    assert forbidden.group(1) == r"[^\x21-\x7E]", forbidden.group(1)
    assert bulk_export_service._PASSWORD_FORBIDDEN_RE.pattern == r"[^\x21-\x7E]"


def test_frontend_disables_execute_button_on_password_error():
    """密码违规时执行按钮不可用（Req 8.4）—— 只给提示不拦按钮等于没拦。"""
    src = _DIALOG.read_bytes().decode("utf-8")
    # canExecute 里必须有 passwordError 这一支
    can_execute = src.split("const canExecute = computed(")[1].split("})")[0]
    assert "passwordError" in can_execute, can_execute


# ═══ 8.4 循环清单由后端下发 ══════════════════════════════════════════════════


def test_cycle_options_cover_dashboard_cycles_and_flag_unsupported():
    """循环清单 = 驾驶舱 D~N ∪ catalog 有 I/E 的循环；不支持的带原因（Req 8.5）。

    🔴 用**真 catalog** 而不是替身：这条判据的价值恰在于「catalog 现在是什么样」——
    E、J 一张可导入导出的 Tab 都没有，正是真栈那个问题的根。用假 catalog 就测不到它。
    """
    from app.services.bulk_tab.manifest_builder import supported_cycles
    from app.services.bulk_tab.scenario_registry import (
        UNSUPPORTED_CYCLE_REASON,
        cycle_options_for_ui,
    )
    from app.services.dashboard_aggregator_service import CYCLES, CYCLE_NAMES

    options = cycle_options_for_ui()
    codes = [o["code"] for o in options]

    # 驾驶舱的每个循环都必须出现（哪怕不支持）—— 只列 catalog 有的会让 E/J 直接消失，
    # 用户仍不知道为什么导不出来
    assert set(CYCLES) <= set(codes)
    assert codes == sorted(codes), "按循环代码升序"
    assert len(codes) == len(set(codes)), "不得重复"

    enabled = set(supported_cycles())
    for opt in options:
        assert opt["name"] == CYCLE_NAMES.get(opt["code"], opt["code"])
        assert opt["supported"] is (opt["code"] in enabled)
        if opt["supported"]:
            assert opt["unsupportedReason"] == ""
        else:
            assert opt["unsupportedReason"] == UNSUPPORTED_CYCLE_REASON

    # 至少有一个支持、一个不支持 —— 否则本判据在两种极端下都恒绿
    assert any(o["supported"] for o in options), "真 catalog 里一个可导入导出的循环都没有？"
    assert any(not o["supported"] for o in options), (
        "真 catalog 里所有 D~N 循环都已接入导入导出 —— "
        "若确已补齐 E/J，请改这条断言并同步 spec，不要直接删掉它"
    )


def test_current_catalog_has_no_import_export_for_e_and_j():
    """现状锚点：E、J 在 catalog 里没有可导入导出的 Tab（Req 8.5 / design §8.4）。

    这是**现状记录而非目标**：spec 明确不在本轮为 E、J 登记 catalog（数据位置不一致，
    E1 数据挂在父底稿上而 bulk 按 `sheet_code` 解析到子底稿，直接登记会得到空数据包）。
    哪天补齐了，本条会打红并提示同步上面那条与 spec。
    """
    from app.services.bulk_tab.manifest_builder import supported_cycles

    enabled = set(supported_cycles())
    assert "E" not in enabled and "J" not in enabled, (
        f"E/J 已出现在可导入导出循环里（现为 {sorted(enabled)}）——"
        "若已补登 catalog，请同步 design §8.4 的「不在本 spec 为 E、J 登记」结论"
    )


def test_unsupported_cycles_reaches_manifest_json():
    """`unsupported_cycles` 进 `manifest.json`（Req 8.5）。"""
    m = _manifest(
        files=[_entry("D2-1")], mode="template", cycles=["D", "E"], unsupported=["E"],
    )
    payload = m.to_dict()
    assert payload["unsupported_cycles"] == ["E"]
    # 导入侧只读 files / cycles / mode，多一个键不影响它们
    assert payload["cycles"] == ["D", "E"] and payload["mode"] == "template"


@pytest.mark.asyncio
async def test_manifest_json_in_zip_carries_unsupported_cycles():
    """ZIP 里的 manifest.json 真的带上了该键（不只是 dataclass 层面）。"""
    m = _manifest(files=[_entry("D2-1")], mode="template", cycles=["D", "E"], unsupported=["E"])
    with _world(m):
        result = await export(
            db=_quiet_db(), project_id=uuid.uuid4(), cycles=["D", "E"], mode="template",
        )
    with zipfile.ZipFile(result) as zf:
        content = json.loads(zf.read("manifest.json"))
        readme = zf.read("README.txt").decode("utf-8")
    assert content["unsupported_cycles"] == ["E"]
    assert "未纳入的循环" in readme and "货币资金" in readme


def test_frontend_cycle_option_fields_are_subset_of_backend_keys():
    """前端 `BulkCycleOption` 的字段 ⊆ 后端下发的键（Req 8.5）。

    前端多声明一个字段就会在运行时恒为 `undefined` —— TS 不会报错（后端响应是 any），
    界面上表现为「原因一栏永远空白」这种没有任何信号的静默失效。
    """
    from app.services.bulk_tab.scenario_registry import cycle_options_for_ui

    src = _DIALOG.read_bytes().decode("utf-8")
    block = src.split("interface BulkCycleOption {")[1].split("}")[0]
    declared = set(re.findall(r"^\s*(\w+)\s*:", block, re.MULTILINE))
    assert declared, block

    backend_keys = set(cycle_options_for_ui()[0])
    assert declared <= backend_keys, f"前端多声明了 {sorted(declared - backend_keys)}"
    # 四个键一个都不能少，否则置灰 / 原因显示会缺料
    assert declared == {"code", "name", "supported", "unsupportedReason"}


# ═══ 端点层：真发请求 ════════════════════════════════════════════════════════


def _bulk_app():
    """只挂 bulk 路由的最小 app。

    🔴 依赖覆盖只 override **内层**的 `get_db` / `get_current_user`，不 override
    `require_project_access("readonly")` —— 那是个依赖**工厂**，每次调用返回新的函数对象，
    拿它当 `dependency_overrides` 的 key 匹配不上，override 会静默失效（全部 401）。
    让真实的鉴权判定跑，只是把它依赖的用户与会话换成替身。
    """
    from fastapi import FastAPI, HTTPException
    from app.core.database import get_db
    from app.deps import get_current_user
    from app.middleware.error_handler import http_exception_handler
    from app.models.base import UserRole
    from app.models.core import ProjectStatus
    from app.routers import wp_bulk_router

    app = FastAPI()
    app.add_exception_handler(HTTPException, http_exception_handler)  # 与 main.py 同一处理器
    app.include_router(wp_bulk_router.router)

    user = SimpleNamespace(id=uuid.uuid4(), username="admin", role=UserRole.admin, is_active=True)
    project = SimpleNamespace(
        name="测试项目", short_name=None, audit_year=2025, status=ProjectStatus.execution,
        wizard_state={},
    )
    db = MagicMock()
    db.get = AsyncMock(return_value=project)
    db.get_bind = MagicMock(return_value=SimpleNamespace(dialect=SimpleNamespace(name="sqlite")))

    async def _db():
        yield db

    app.dependency_overrides[get_db] = _db
    app.dependency_overrides[get_current_user] = lambda: user
    return app


@pytest.mark.asyncio
async def test_scenarios_endpoint_returns_cycle_list():
    """场景端点下发 `cycles`（Req 8.5）—— 前端据此渲染，不再写死 D~N。"""
    from httpx import ASGITransport, AsyncClient

    app = _bulk_app()
    async with AsyncClient(transport=ASGITransport(app=app), base_url="http://test") as client:
        resp = await client.get(f"/api/projects/{uuid.uuid4()}/bulk-tab/scenarios")

    assert resp.status_code == 200, resp.text
    payload = resp.json()
    body = payload.get("data", payload)
    assert "scenarios" in body and body["scenarios"], "场景列表不应为空"
    cycles = body.get("cycles")
    assert isinstance(cycles, list) and cycles, "端点未下发 cycles"
    assert set(cycles[0]) == {"code", "name", "supported", "unsupportedReason"}
    assert any(not c["supported"] for c in cycles), "真 catalog 下应有不支持的循环（E/J）"


@pytest.mark.asyncio
@pytest.mark.parametrize("endpoint", ["export-templates", "export-data"])
async def test_endpoint_returns_422_json_with_chinese_reason_when_nothing_to_export(endpoint):
    """零 Tab ⇒ 端点返回 422 JSON 中文原因，**不是** ZIP（Req 8.1）。

    选 422 而非 4xx 其它码 / 5xx 的理由（design §8.0）：这是用户可纠正的请求问题；
    且前端拦截器对 5xx 会自动重试两遍 —— 零 Tab 重跑两遍毫无意义，只让用户多等十分钟。
    """
    from httpx import ASGITransport, AsyncClient

    app = _bulk_app()
    manifest = _manifest(files=[_entry("D2-1")], skipped=[])
    with _world(manifest, fail_tabs=True), \
            patch("app.routers.wp_bulk_router.make_bulk_visible_filter", lambda *a, **k: _allow_all):
        async with AsyncClient(transport=ASGITransport(app=app), base_url="http://test") as client:
            resp = await client.post(
                f"/api/projects/{uuid.uuid4()}/bulk-tab/{endpoint}",
                json={"cycles": ["D"]},
            )

    assert resp.status_code == 422, resp.text
    assert resp.headers["content-type"].startswith("application/json")
    body = resp.json()
    assert body["code"] == 422
    assert "没有可导出的底稿" in body["message"]
    # 中文原因里不得夹带英文 skip_reason 代码
    for code in SKIP_REASON_LABELS:
        assert code not in body["message"]


async def _allow_all(wp_id, sheet_code):  # noqa: ANN001, ANN201
    return True


def test_skip_reason_label_and_advice_tables_cover_the_same_codes():
    """两张文案表覆盖同一组代码 —— 漏一个就会在 README 或拦截提示里露出英文代码。"""
    assert set(SKIP_REASON_LABELS) == set(SKIP_REASON_ADVICE)
    # 生产代码里真实会产生的 skip_reason 全部在表内
    produced = {"no_adapter", "export_failed", "no_data", "unchanged", "resolve_instance_miss", "unknown"}
    assert produced <= set(SKIP_REASON_LABELS), sorted(produced - set(SKIP_REASON_LABELS))


@pytest.mark.asyncio
async def test_index_sheet_status_column_is_chinese():
    """`_底稿目录.xlsx` 的状态列用中文，且整循环未纳入的也在目录里（Req 8.2 / 8.5）。

    🔴 这条必须真读 xlsx 单元格：目录表是二进制产物，只断言「ZIP 里有这个文件」
    对状态列文案是恒绿的 —— 退回英文 `跳过(no_adapter)` 照样通过。
    """
    import openpyxl

    from app.services.bulk_tab.scenario_registry import UNSUPPORTED_CYCLE_REASON

    m = _manifest(
        files=[_entry("D2-1")],
        skipped=[_skipped("D2-8", "no_adapter"), _skipped("D2-9", "resolve_instance_miss")],
        unsupported=["E"],
        mode="template",
        cycles=["D", "E"],
    )
    with _world(m):
        result = await export(
            db=_quiet_db(), project_id=uuid.uuid4(), cycles=["D", "E"], mode="template",
        )

    with zipfile.ZipFile(result) as zf:
        raw = zf.read("_底稿目录.xlsx")
    ws = openpyxl.load_workbook(io.BytesIO(raw), read_only=True).active
    rows = [[c for c in row] for row in ws.iter_rows(values_only=True)]
    header, *body = rows
    assert header[4] == "状态"

    by_code = {r[2]: r for r in body if r[2]}
    assert by_code["D2-1"][4] == "已导出"
    for code, reason in (("D2-8", "no_adapter"), ("D2-9", "resolve_instance_miss")):
        status = by_code[code][4]
        assert status == f"未导出（{SKIP_REASON_LABELS[reason]}）", status
        assert reason not in status, "状态列不得出现英文 skip_reason 代码"

    # 整个循环未纳入的也要在目录里，否则用户勾了 E 却在目录中查无此项
    cycle_rows = [r for r in body if r[1] == "E" and not r[2]]
    assert cycle_rows, "unsupported_cycles 未出现在底稿目录里"
    assert UNSUPPORTED_CYCLE_REASON in str(cycle_rows[0][4])


# ═══ build_manifest 真实调用（不 mock）═════════════════════════════════════════
#
# 🔴 补这一节的原因：上面所有用例都用 `_world()` 把 `build_manifest` 换成替身，于是
# 「不支持的循环进 unsupported_cycles」这条逻辑**从未被真正执行过** —— 变异检验里把
# `manifest.unsupported_cycles.append(cycle_code)` 换成 `pass` 时 40 个用例全绿。
# 判据必须至少有一处走真实实现。


@pytest.mark.asyncio
async def test_build_manifest_records_cycle_with_no_import_export_tabs():
    """真 catalog + 真 build_manifest：E 循环无 I/E Tab ⇒ 记入 unsupported_cycles。"""
    from app.services.bulk_tab.manifest_builder import build_manifest

    db = AsyncMock()
    manifest = await build_manifest(
        db=db, project_id=uuid.uuid4(), cycles=["E"], mode="template", audit_year=2025,
    )

    assert manifest.unsupported_cycles == ["E"]
    assert manifest.files == [] and manifest.skipped == []
    # 不支持的循环应在两次 DB 查询之前就短路（catalog 为空说明再查也是空）
    db.execute.assert_not_awaited()


@pytest.mark.asyncio
async def test_build_manifest_does_not_flag_supported_cycle():
    """正向对照：catalog 里有 I/E Tab 的循环**不**进 unsupported_cycles。

    没有这条，把 `if not _cycle_sheets` 改成 `if True` 也会让上一条绿
    —— 那等于所有循环都被判成「未接入」，功能全废。
    """
    from app.services.bulk_tab.manifest_builder import build_manifest

    with patch("app.services.wp_bulk_tab_export.list_export_sheets", AsyncMock(return_value=[])), \
            patch("app.services.acnr.manifest.list_import_export", AsyncMock(return_value=[])):
        manifest = await build_manifest(
            db=AsyncMock(), project_id=uuid.uuid4(), cycles=["D"], mode="template", audit_year=2025,
        )

    assert manifest.unsupported_cycles == []
    assert manifest.cycles == ["D"]


@pytest.mark.asyncio
async def test_build_manifest_default_scope_is_supported_cycles_only():
    """`cycles=None` 时范围 = catalog 有 I/E 的循环 ⇒ 默认范围里不该有 unsupported。"""
    from app.services.bulk_tab.manifest_builder import build_manifest, supported_cycles

    with patch("app.services.wp_bulk_tab_export.list_export_sheets", AsyncMock(return_value=[])), \
            patch("app.services.acnr.manifest.list_import_export", AsyncMock(return_value=[])):
        manifest = await build_manifest(
            db=AsyncMock(), project_id=uuid.uuid4(), cycles=None, mode="template", audit_year=2025,
        )

    assert manifest.cycles == supported_cycles()
    assert manifest.unsupported_cycles == []

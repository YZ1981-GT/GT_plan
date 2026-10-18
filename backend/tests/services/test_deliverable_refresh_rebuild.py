"""交付件增量刷新按模板重建块 + 交付过滤（spec ``deliverable-trio-generation-accuracy`` §16b）。

覆盖：
- 模板结构块（标题样式 + 已渲染表）走 ``rebuild_block``：库内表重渲染、大模型草稿 / 指引被交付过滤剔除；
- 纯文字块回退仍接入交付过滤（Task 9 口径）；
- anchor / marker 两定位模式都正确；
- ``_replace_block_content`` 返回的 Rendered_Block_Hash 口径不变（与导出 / 人工编辑检测侧同一函数）。

磁盘变异（在被测函数的下一层注入，不替换生产函数本身）各自打红：
- 关掉「模板重建分流」（``_has_rebuildable_structure`` 恒 False）→ 库内表不再重渲染，模板重建用例红；
- 关掉「交付过滤接入」（``make_block_filter`` 退化为恒等）→ 大模型草稿残留，过滤用例红。
"""

from __future__ import annotations

import asyncio
import uuid
from io import BytesIO

import pytest
from docx import Document
from docx.oxml.ns import qn
from sqlalchemy.dialects.sqlite.base import SQLiteTypeCompiler
from sqlalchemy.ext.asyncio import async_sessionmaker, create_async_engine

from app.models.audit_platform_models import DeliverableSectionState
from app.models.phase13_models import WordExportTask
from app.models.report_models import DisclosureNote
from app.services.deliverable_refresh_service import DeliverableRefreshService
from app.services.section_anchor_utils import (
    SectionBlock,
    block_text_of,
    resolve_section_blocks,
    scan_anchor_blocks,
    write_section_anchors,
)
from app.services.word_doc_utils import remove_section_markers, scan_section_blocks

SQLiteTypeCompiler.visit_JSONB = SQLiteTypeCompiler.visit_JSON

PROJECT_ID = uuid.uuid4()
TASK_ID = uuid.uuid4()
ACTOR_ID = uuid.uuid4()
YEAR = 2025

_W_TBL = qn("w:tbl")
_W_T = qn("w:t")

_TABLES = [
    DeliverableSectionState.__table__,
    DisclosureNote.__table__,
    WordExportTask.__table__,
]

#: 库内附注（刷新源）：新表数据 + 正文含「本公司…」正式陈述 + 一行大模型提示（应被过滤）
NEW_TABLE = {
    "_tables": [
        {
            "name": "货币资金",
            "headers": ["项目", "期末余额", "期初余额"],
            "rows": [
                {"label": "库存现金", "values": ["库存现金", 500.0, 300.0]},
                {"label": "银行存款", "values": ["银行存款", 9000.0, 8000.0]},
            ],
        }
    ]
}
NEW_TEXT = "本公司货币资金主要为银行存款。\n审计师提示：请提供各明细科目的具体金额。"


def _bytes(doc) -> bytes:
    buf = BytesIO()
    doc.save(buf)
    return buf.getvalue()


def _template_block_doc(mode: str) -> bytes:
    """交付件形态：标题样式段 + 表题 + 一张旧渲染表 + 旧正文。mode='anchor' 清标记写锚点。"""
    doc = Document()
    doc.add_paragraph("首段")
    doc.add_paragraph("##SECTION:八、1##")
    doc.add_paragraph("八、1 货币资金", style="Heading 2")
    doc.add_paragraph("货币资金")  # 表题（caption-like）
    tbl = doc.add_table(rows=2, cols=3)
    for i, h in enumerate(["项目", "期末余额", "期初余额"]):
        tbl.cell(0, i).text = h
    tbl.cell(1, 0).text = "库存现金"
    tbl.cell(1, 1).text = "100.00"
    tbl.cell(1, 2).text = "90.00"
    doc.add_paragraph("旧正文一段。")
    doc.add_paragraph("##/SECTION:八、1##")
    if mode == "anchor":
        blocks = {b.section_code: b for b in scan_section_blocks(doc)}
        write_section_anchors(
            doc,
            [
                SectionBlock(section_code=c, open_el=b.open_el, close_el=b.close_el)
                for c, b in blocks.items()
            ],
        )
        remove_section_markers(doc)
    return _bytes(doc)


def _plain_block_doc(mode: str) -> bytes:
    """纯文字块（无标题样式 / 无表）：只有正文段。"""
    doc = Document()
    doc.add_paragraph("首段")
    doc.add_paragraph("##SECTION:九、1##")
    doc.add_paragraph("旧正文。")
    doc.add_paragraph("##/SECTION:九、1##")
    if mode == "anchor":
        blocks = {b.section_code: b for b in scan_section_blocks(doc)}
        write_section_anchors(
            doc,
            [
                SectionBlock(section_code=c, open_el=b.open_el, close_el=b.close_el)
                for c, b in blocks.items()
            ],
        )
        remove_section_markers(doc)
    return _bytes(doc)


def _run(scenario):
    async def _main():
        engine = create_async_engine("sqlite+aiosqlite:///:memory:", echo=False)
        async with engine.begin() as conn:
            await conn.run_sync(
                DeliverableSectionState.metadata.create_all, tables=_TABLES
            )
        factory = async_sessionmaker(engine, expire_on_commit=False)
        try:
            async with factory() as session:
                return await scenario(session)
        finally:
            await engine.dispose()

    return asyncio.run(_main())


async def _setup(session, *, section, text, table_data):
    session.add(
        WordExportTask(
            id=TASK_ID,
            project_id=PROJECT_ID,
            doc_type="disclosure_notes",
            status="editing",
            created_by=ACTOR_ID,
        )
    )
    session.add(
        DisclosureNote(
            project_id=PROJECT_ID,
            year=YEAR,
            note_section=section,
            section_title=section,
            text_content=text,
            table_data=table_data,
            is_deleted=False,
        )
    )
    await session.flush()


def _refresh(docx_bytes, section, *, setup_text, setup_table):
    """跑单章节刷新（mock render_and_store 只取新字节），返回刷新后的 docx bytes。"""

    async def _scenario(session):
        await _setup(session, section=section, text=setup_text, table_data=setup_table)
        svc = DeliverableRefreshService(session)
        stored: dict = {}

        async def _fake_store(_self, task_id, **kw):
            stored["bytes"] = kw["docx_bytes"]

            class _V:
                version_no = 2

            class _R:
                version = _V()

            return _R()

        from app.services import deliverable_service as ds_mod

        original = ds_mod.DeliverableService.render_and_store
        ds_mod.DeliverableService.render_and_store = _fake_store  # type: ignore[assignment]
        try:
            result = await svc.refresh_section(
                TASK_ID, PROJECT_ID, YEAR, section, ACTOR_ID, docx_bytes=docx_bytes
            )
        finally:
            ds_mod.DeliverableService.render_and_store = original  # type: ignore[assignment]
        return result, stored.get("bytes")

    return _run(_scenario)


def _block_tables(docx_bytes, section):
    doc = Document(BytesIO(docx_bytes))
    for b in scan_anchor_blocks(doc) or []:
        if b.section_code == section:
            return [e for e in b.elements if e.tag == _W_TBL]
    # marker 模式
    for b in scan_section_blocks(doc):
        if b.section_code == section:
            return [e for e in b.elements if e.tag == _W_TBL]
    return []


def _table_cell_texts(table) -> list[str]:
    out = []
    for tr in table.findall(qn("w:tr")):
        for tc in tr.findall(qn("w:tc")):
            out.append("".join(t.text or "" for t in tc.iter(_W_T)))
    return out


def _block_text(docx_bytes, section) -> str:
    doc = Document(BytesIO(docx_bytes))
    for b in scan_anchor_blocks(doc) or []:
        if b.section_code == section:
            return block_text_of(b.elements)
    for b in scan_section_blocks(doc):
        if b.section_code == section:
            return block_text_of(b.elements)
    return ""


# ─── 模板结构块：重渲染库内表 + 过滤正文 ──────────────────────────────────────


@pytest.mark.parametrize("mode", ["anchor", "marker"])
def test_template_block_rerenders_table_and_filters_text(mode):
    docx = _template_block_doc(mode)
    original_mode = resolve_section_blocks(Document(BytesIO(docx)))[0]
    assert original_mode == mode

    result, new_bytes = _refresh(
        docx, "八、1", setup_text=NEW_TEXT, setup_table=NEW_TABLE
    )
    assert result["refreshed"] == ["八、1"], result
    assert new_bytes is not None

    # 1. 库内表重渲染：出现新数据（9000.00 / 银行存款 的新期初 8000.00），旧值 100.00 不在
    tables = _block_tables(new_bytes, "八、1")
    assert tables, "重建后块内应仍有表格"
    cells = _table_cell_texts(tables[0])
    joined = " ".join(cells)
    assert "银行存款" in joined
    assert "9,000.00" in joined or "9000.00" in joined, joined
    assert "100.00" not in joined, "旧表数据应被重渲染替换"

    # 2. 交付过滤：正式陈述保留，大模型提示被剔除
    text = _block_text(new_bytes, "八、1")
    assert "本公司货币资金主要为银行存款。" in text
    assert "审计师提示" not in text, text


# ─── 纯文字块：回退仍接交付过滤 ───────────────────────────────────────────────


@pytest.mark.parametrize("mode", ["anchor", "marker"])
def test_plain_block_fallback_applies_delivery_filter(mode):
    docx = _plain_block_doc(mode)
    result, new_bytes = _refresh(
        docx,
        "九、1",
        setup_text="本年度无重大或有事项。\n请补充相关合同金额并替换本段。",
        setup_table=None,
    )
    assert result["refreshed"] == ["九、1"], result
    text = _block_text(new_bytes, "九、1")
    assert "本年度无重大或有事项。" in text
    assert "请补充" not in text, text
    assert "旧正文。" not in text


# ─── Rendered_Block_Hash 口径不变 ────────────────────────────────────────────


def test_rendered_block_hash_matches_filtered_content():
    """刷新后 rendered_block_hash = 过滤后块内文字的规范化 sha256（与 block_text_hash 同口径）。"""
    import hashlib

    from app.services.section_anchor_utils import normalize_block_text

    docx = _template_block_doc("anchor")

    async def _scenario(session):
        await _setup(session, section="八、1", text=NEW_TEXT, table_data=NEW_TABLE)
        # 写基线行（使 clear_section_stale 有行可更新）
        session.add(
            DeliverableSectionState(
                word_export_task_id=TASK_ID,
                project_id=PROJECT_ID,
                year=YEAR,
                section_code="八、1",
                source_snapshot_hash="0" * 64,
                rendered_block_hash=None,
                is_stale=True,
            )
        )
        await session.flush()
        svc = DeliverableRefreshService(session)
        stored: dict = {}

        async def _fake_store(_self, task_id, **kw):
            stored["bytes"] = kw["docx_bytes"]

            class _V:
                version_no = 2

            class _R:
                version = _V()

            return _R()

        from app.services import deliverable_service as ds_mod

        original = ds_mod.DeliverableService.render_and_store
        ds_mod.DeliverableService.render_and_store = _fake_store  # type: ignore[assignment]
        try:
            await svc.refresh_section(
                TASK_ID, PROJECT_ID, YEAR, "八、1", ACTOR_ID, docx_bytes=docx
            )
        finally:
            ds_mod.DeliverableService.render_and_store = original  # type: ignore[assignment]

        import sqlalchemy as sa

        row = (
            await session.execute(
                sa.select(DeliverableSectionState).where(
                    DeliverableSectionState.section_code == "八、1"
                )
            )
        ).scalar_one()
        return row.rendered_block_hash, row.is_stale, stored["bytes"]

    stored_hash, is_stale, new_bytes = _run(_scenario)
    assert is_stale is False

    # 期望 = 刷新后块内文字（表格不计入、## 行剔除）规范化后的 sha256
    text = _block_text(new_bytes, "八、1")
    expected = hashlib.sha256(
        normalize_block_text(text).encode("utf-8")
    ).hexdigest()
    assert stored_hash == expected, (stored_hash, expected, text)
    # 内容里确实只剩过滤后的正文 + 标题（提示已删）
    assert "审计师提示" not in text

# ─── 加固 1：真实交付件 w:sdt 内容控件路径（复盘补）──────────────────────────
#
# 真实交付件把章节内部内容包进 w:sdt 内容控件（note_word_exporter step6.5 调
# inject_content_controls_for_blocks，默认开）。此时 anchor_block_content_host 走
# sdtContent 分支、哨兵要插进 sdtContent 内部。现有 anchor 用例内容在 body 级，
# 不覆盖这条真实路径 —— 用生产注入器造夹具（不臆造 XML）。


def _sdt_wrapped_template_doc() -> bytes:
    """交付件真实形态：章节内部内容被 inject_content_controls_for_blocks 包进 w:sdt，
    再写锚点、清 ##SECTION## 标记。与全量导出 step6.5 同一生产路径。"""
    from app.services.content_control_injector import (
        inject_content_controls_for_blocks,
    )

    doc = Document()
    doc.add_paragraph("首段")
    doc.add_paragraph("##SECTION:八、1##")
    doc.add_paragraph("八、1 货币资金", style="Heading 2")
    doc.add_paragraph("货币资金")
    tbl = doc.add_table(rows=2, cols=3)
    for i, h in enumerate(["项目", "期末余额", "期初余额"]):
        tbl.cell(0, i).text = h
    tbl.cell(1, 0).text = "库存现金"
    tbl.cell(1, 1).text = "100.00"
    tbl.cell(1, 2).text = "90.00"
    doc.add_paragraph("旧正文一段。")
    doc.add_paragraph("##/SECTION:八、1##")
    # 先按标记注入内容控件（只包内部内容，开闭标记留 body 级）
    n = inject_content_controls_for_blocks(doc, scan_section_blocks(doc))
    assert n == 1, "夹具前置：应注入 1 个内容控件"
    # 再写锚点、清标记（交付件最终形态）
    blocks = {b.section_code: b for b in scan_section_blocks(doc)}
    write_section_anchors(
        doc,
        [
            SectionBlock(section_code=c, open_el=b.open_el, close_el=b.close_el)
            for c, b in blocks.items()
        ],
    )
    remove_section_markers(doc)
    return _bytes(doc)


def _sdt_count(docx_bytes) -> int:
    doc = Document(BytesIO(docx_bytes))
    return len(doc.element.body.findall(".//" + qn("w:sdt")))


def _content_inside_sdt(docx_bytes, section) -> bool:
    """该章节的内容元素是否仍在 w:sdtContent 内（没被插到控件外）。"""
    doc = Document(BytesIO(docx_bytes))
    sdt_content = qn("w:sdtContent")
    for b in scan_anchor_blocks(doc) or []:
        if b.section_code != section:
            continue
        # anchor_block_content_host 恰有一个 sdt 时返回 sdtContent 宿主
        from app.services.section_anchor_utils import anchor_block_content_host

        host, children = anchor_block_content_host(b)
        return host.tag == sdt_content and len(children) > 0
    return False


def test_sdt_content_control_path_rebuilds_inside_control():
    """内容被 w:sdt 包裹的真实交付件：刷新后库内表重渲染进 sdtContent 内、过滤生效、
    哨兵清理干净、内容仍在控件内（不落控件外）。"""
    docx = _sdt_wrapped_template_doc()
    assert _sdt_count(docx) == 1, "夹具应有一个内容控件"
    # 定位模式必是 anchor（标记已清）
    assert resolve_section_blocks(Document(BytesIO(docx)))[0] == "anchor"

    result, new_bytes = _refresh(
        docx, "八、1", setup_text=NEW_TEXT, setup_table=NEW_TABLE
    )
    assert result["refreshed"] == ["八、1"], result

    # 1. 内容控件仍在、未增殖（哨兵是临时 w:p，不应变成新 sdt）
    assert _sdt_count(new_bytes) == 1

    # 2. 库内表重渲染进了控件内部
    assert _content_inside_sdt(new_bytes, "八、1"), "刷新后内容应仍在 sdtContent 内"
    tables = _block_tables(new_bytes, "八、1")
    assert tables, "重建后块内应有表"
    joined = " ".join(_table_cell_texts(tables[0]))
    assert "银行存款" in joined and "100.00" not in joined, joined

    # 3. 过滤生效、哨兵不残留（控件内无多余空段 —— 正文段都有文字或是表）
    text = _block_text(new_bytes, "八、1")
    assert "本公司货币资金主要为银行存款。" in text
    assert "审计师提示" not in text, text


# ─── 加固 2：_has_rebuildable_structure 分流边界的反向变异（复盘补）──────────


def test_table_block_degrades_to_fallback_when_tbl_predicate_killed(monkeypatch):
    """反向变异：把 _has_rebuildable_structure 的 w:tbl 判据改成恒不认表，
    有表的块应从「模板重建」退化到「纯文字回退」—— 库内表不再出现在块内。
    变异注在被测函数的下一层（patch 模块内的 _heading_level 不动，单独 patch 判据里的
    tbl 识别），不替换 rebuild_refresh_block 本身。"""
    import app.services.deliverable_refresh_rebuild as rr

    # 先确认未变异时是重建路径（有表）
    docx = _template_block_doc("anchor")
    _, baseline_bytes = _refresh(
        docx, "八、1", setup_text=NEW_TEXT, setup_table=NEW_TABLE
    )
    assert _block_tables(baseline_bytes, "八、1"), "未变异：应走重建、块内有表"

    # 变异：让分流判据看不到任何 w:tbl（只保留标题样式判据）——但本夹具标题是
    # Heading 2，仍会走重建；为隔离「表判据」，构造一个只有表、无标题样式的块。
    orig = rr._has_rebuildable_structure

    def _no_tbl(adapted, names):
        # 下一层变异：屏蔽表识别，等价于 w:tbl 判据失效
        from app.services.note_template_docx_base import _heading_level

        inner = adapted.elements[1:-1]
        for el in inner:
            if el.tag == rr._W_P and _heading_level(el, names):
                return True
        return False

    monkeypatch.setattr(rr, "_has_rebuildable_structure", _no_tbl)
    assert rr._has_rebuildable_structure is not orig

    # 用「只有表、无标题样式」的块：变异后应退化到回退（不重渲染表）
    doc = Document()
    doc.add_paragraph("##SECTION:八、1##")
    t = doc.add_table(rows=2, cols=3)
    for i, h in enumerate(["项目", "期末余额", "期初余额"]):
        t.cell(0, i).text = h
    t.cell(1, 0).text = "库存现金"
    t.cell(1, 1).text = "100.00"
    doc.add_paragraph("旧正文。")
    doc.add_paragraph("##/SECTION:八、1##")
    blocks = {b.section_code: b for b in scan_section_blocks(doc)}
    write_section_anchors(
        doc,
        [SectionBlock(section_code=c, open_el=b.open_el, close_el=b.close_el)
         for c, b in blocks.items()],
    )
    remove_section_markers(doc)
    _, mutated_bytes = _refresh(
        _bytes(doc), "八、1", setup_text=NEW_TEXT, setup_table=NEW_TABLE
    )
    # 变异后走回退：回退只渲染过滤后正文，不重渲染库内表 → 块内无库内表数据
    joined = " ".join(
        _table_cell_texts(tb) for tb in _block_tables(mutated_bytes, "八、1")
    ) if _block_tables(mutated_bytes, "八、1") else ""
    assert "9,000.00" not in joined and "9000.00" not in joined, (
        "变异屏蔽表判据后，有表块应退化到回退、不应重渲染库内表", joined
    )


def test_unmapped_heading_style_falls_back_not_crash():
    """标题样式名不在 style_map（_heading_level 返 None）时：分流到纯文字回退，
    不崩溃；库内正文经过滤后写入。"""
    doc = Document()
    doc.add_paragraph("##SECTION:九、2##")
    # 自定义/未映射样式名的「标题」段：_heading_level 认不出 → 不算标题结构
    p = doc.add_paragraph("九、2 自定义标题")
    p.style = doc.styles["Normal"]  # 普通样式，无 heading 级别
    doc.add_paragraph("旧正文。")
    doc.add_paragraph("##/SECTION:九、2##")
    blocks = {b.section_code: b for b in scan_section_blocks(doc)}
    write_section_anchors(
        doc,
        [SectionBlock(section_code=c, open_el=b.open_el, close_el=b.close_el)
         for c, b in blocks.items()],
    )
    remove_section_markers(doc)

    result, new_bytes = _refresh(
        _bytes(doc),
        "九、2",
        setup_text="本年度无重大事项。\n请补充金额。",
        setup_table=None,
    )
    assert result["refreshed"] == ["九、2"], result
    text = _block_text(new_bytes, "九、2")
    assert "本年度无重大事项。" in text
    assert "请补充" not in text, text
    assert "旧正文。" not in text

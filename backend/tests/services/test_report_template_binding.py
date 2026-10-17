"""报表模板行绑定（spec deliverable-trio-generation-accuracy §1.1）。

库内 report_config 与模板占位是两套编号（真实库 soe_standalone 226 个占位 113 个同码不同名），
旧导出按编码直取 ⇒ 重药安徽「衍生金融资产」格拿到应收账款 5.96 亿。这里钉住：
按行名绑定、同码不同名不取数、同名多份取报表合计引用的那份、模板同码多处取「本位」。
"""

from __future__ import annotations

from decimal import Decimal

import pytest
from openpyxl import Workbook

from app.services.report_template_binding import (
    best_rows_by_key,
    bind_rows_to_template,
    load_template_slots,
    normalize_row_label,
    period_columns,
    referenced_row_codes,
    scan_workbook_slots,
)


@pytest.mark.parametrize(
    ("text", "expected"),
    [
        ("  货币资金", ("货币资金", "")),
        ("△结算备付金", ("结算备付金", "△")),
        ("  *2.少数股东损益", ("少数股东损益", "*")),
        ("    其中：应收股利", ("应收股利", "")),
        ("         加： 营业外收入", ("营业外收入", "")),
        ("四、利润总额（亏损总额以“－”号填列）", ("利润总额", "")),
        ("投资收益（损失以“-”号填列）", ("投资收益", "")),
        ("实收资本（或股本）", ("实收资本", "")),
        ("（一）按所有权归属分类", ("按所有权归属分类", "")),
        ("上期期末余额", ("上年年末余额", "")),
        ("#减：已归还投资", ("已归还投资", "#")),
        ("1. 重新计量设定受益计划变动额", ("重新计量设定受益计划变动额", "")),
        ("负 债 合 计", ("负债合计", "")),
    ],
)
def test_normalize_row_label(text, expected):
    assert normalize_row_label(text) == expected


def _template(rows: list[tuple[str, str]], title: str = "资产负债表"):
    """合成模板：(行名, 编码或 =公式)；第 4 行表头「期末余额 / 期初余额」。"""
    wb = Workbook()
    ws = wb.active
    ws.title = title
    ws["A4"], ws["C4"], ws["D4"] = "项目", "期末余额", "期初余额"
    for i, (label, code) in enumerate(rows, start=5):
        if label:
            ws.cell(i, 1, label)
        if code.startswith("="):
            ws.cell(i, 3, code)
        else:
            ws.cell(i, 3, f"{{{{row:{code}:current}}}}")
            ws.cell(i, 4, f"{{{{row:{code}:prior}}}}")
    return wb


def _row(code: str, name: str, amount=None, *, formula: str | None = None, prior=None) -> dict:
    return {
        "row_code": code,
        "row_name": name,
        "current_period_amount": None if amount is None else Decimal(str(amount)),
        "prior_period_amount": None if prior is None else Decimal(str(prior)),
        "formula_used": formula,
    }


def test_code_drift_binds_by_name_and_never_by_code():
    """库内 BS-006 = 应收账款（通用编号）；模板 BS-006 = 衍生金融资产、BS-008 = 应收账款。"""
    slots = scan_workbook_slots(_template([("货币资金", "BS-002"), ("衍生金融资产", "BS-006"), ("应收账款", "BS-008")]))
    res = bind_rows_to_template(
        [_row("BS-002", "货币资金", 100, formula="TB('1001')"), _row("BS-006", "应收账款", 500, formula="TB('1122')")],
        slots,
    )
    assert res.lookup("资产负债表", 7, "BS-008")["current_period_amount"] == Decimal("500")
    assert res.lookup("资产负债表", 7, "BS-008")["db_row_code"] == "BS-006"
    assert res.lookup("资产负债表", 6, "BS-006") is None  # 同码不同名：不取数
    assert [(s.code, name) for s, name in res.conflicts] == [("BS-006", "应收账款")]


def test_duplicate_name_prefers_row_referenced_by_totals():
    """同名两份（通用含应付利息 / 模板编号不含）：取报表合计引用的那份，导出合计才与报表模块一致。"""
    slots = scan_workbook_slots(_template([("其他应付款", "BS-075"), ("流动负债合计", "=SUM(C5:C5)")]))
    rows = [
        _row("BS-050", "其他应付款", 129, formula="TB('2241') + TB('2231')"),
        _row("BS-054", "流动负债合计", 129, formula="ROW('BS-041') + ROW('BS-050')"),
        _row("BS-075", "其他应付款", 127, formula="TB('2241')"),
    ]
    res = bind_rows_to_template(rows, slots)
    assert res.lookup("资产负债表", 5, "BS-075")["db_row_code"] == "BS-050"
    assert best_rows_by_key(rows)[("BS", "其他应付款", "")]["row_code"] == "BS-050"


def test_weak_same_code_row_yields_to_row_with_formula():
    slots = scan_workbook_slots(_template([("预付款项", "BS-010")]))
    res = bind_rows_to_template(
        [_row("BS-008", "预付款项", 13, formula="TB('1123')"), _row("BS-010", "预付款项", 0)],
        slots,
    )
    assert res.lookup("资产负债表", 5, "BS-010")["db_row_code"] == "BS-008"


def test_same_code_repeated_in_template_goes_to_its_home_position():
    """模板把 BS-111「其中：优先股」写了两处（应付债券下 / 其他权益工具下）：前后编码连号的才是本位。"""
    slots = scan_workbook_slots(_template([
        ("应付债券", "BS-086"), ("其中：优先股", "BS-111"), ("其他权益工具", "BS-110"), ("其中：优先股", "BS-111"),
    ]))
    res = bind_rows_to_template([_row("BS-111", "其中：优先股", 5)], slots)
    assert res.lookup("资产负债表", 8, "BS-111")["current_period_amount"] == Decimal("5")
    assert res.lookup("资产负债表", 6, "BS-111") is None


def test_unlabeled_template_row_falls_back_to_code():
    slots = scan_workbook_slots(_template([("", "BS-002")]))
    res = bind_rows_to_template([_row("BS-002", "货币资金", 7)], slots)
    assert res.lookup("资产负债表", 5, "BS-002")["current_period_amount"] == Decimal("7")
    assert res.methods["code_unlabeled"] == 1


def test_unbound_nonzero_excludes_template_formula_rows():
    slots = scan_workbook_slots(_template([("货币资金", "BS-002"), ("资产总计", "=C5")]))
    res = bind_rows_to_template(
        [_row("BS-002", "货币资金", 1), _row("BS-039", "资产总计", 1), _row("BS-099", "特殊项目", 3)],
        slots,
    )
    assert [r["row_code"] for r in res.unbound_nonzero] == ["BS-099"]


def test_referenced_row_codes_parses_row_and_sum_row_ranges():
    rows = [
        _row("IS-041", "营业外收入"), _row("IS-042", "其中：政府补助"), _row("IS-043", "营业外支出"),
        _row("IS-044", "利润总额", formula="SUM_ROW('IS-041','IS-043')"),
        _row("BS-099", "负债和所有者权益总计", formula="ROW('BS-070') + ROW( 'BS-091' )"),
    ]
    assert referenced_row_codes(rows) == {"IS-041", "IS-042", "IS-043", "BS-070", "BS-091"}


def test_period_columns_by_header_text():
    ws = _template([("货币资金", "BS-002")]).active
    assert period_columns(ws) == (3, 4)


def test_real_soe_template_slots_are_labelled():
    """真实国企单体模板：每个出现位都有行名（编码兜底永不触发），且「实收资本」是模板自算行。"""
    slots = load_template_slots("soe_standalone")
    if slots is None:
        pytest.skip("报表模板缺失")
    assert slots.slots and all(s.key for s in slots.slots)
    receivable = [s for s in slots.slots if s.key == "应收账款" and s.family == "BS"]
    assert [s.code for s in receivable] == ["BS-008"]
    assert ("BS", "实收资本", "") in slots.formula_keys

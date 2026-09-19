"""K 系披露 `sheet_name` 常量 ↔ 源 xlsx tab 名逐字守卫。

**为什么需要这条**：前端 `disclosureSheetNameRegistry.spec.ts` 只比对
「`X_DISCLOSURE_SHEET_NAME` 常量 ↔ `note_workpaper_sync_registry.json`」，而 registry
是**由常量生成**的 → 常量与源 xlsx 的漂移它查不出。2026-07-30 实测 K2 与 K8~K13
共 13 处写成半角括号（源 xlsx 是全角），会让附注「打开同步底稿」的 `?sheet=`
精确匹配落空，跳回底稿首个 sheet。

本测试直接用 openpyxl 读 `wb.sheetnames`，与 `.ts` 源码里的常量字面量比对。

括号写法**原样保留**：源 xlsx 本身混用半/全角（K3 两版都是半角、K5 上市前全角后半角、
K6 国企前半角后全角、K7 国企写「国有企业」），故断言是「与 xlsx 一致」而非「统一全角」。

spec: .kiro/specs/k-cycle-disclosure-alignment/ R1.4（Task 4.2）
"""
from __future__ import annotations

import re
from pathlib import Path

import pytest

try:
    import openpyxl
except ImportError:  # pragma: no cover - CI 环境未装 openpyxl 时跳过
    openpyxl = None  # type: ignore[assignment]

_BACKEND = Path(__file__).resolve().parents[2]
_REPO = _BACKEND.parent
TEMPLATE_DIR = _BACKEND / "wp_templates" / "K"
MAP_DIR = _REPO / "audit-platform" / "frontend" / "src" / "components" / "workpaper" / "composables"

# 本 spec 批 1 收口范围 + 已收口的 K2（防回退）
CYCLES = ["K2", "K8", "K9", "K10", "K11", "K12", "K13"]

_CONST_RE = re.compile(
    r"export const (?P<code>K\d+)_DISCLOSURE_SHEET_NAME\s*=\s*\{(?P<body>.*?)\}",
    re.S,
)


def _read_constants(code: str) -> dict[str, str]:
    path = MAP_DIR / f"{code.lower()}NoteSectionMap.ts"
    if not path.exists():
        pytest.skip(f"{path.name} 不存在")
    src = path.read_text(encoding="utf-8")
    m = _CONST_RE.search(src)
    assert m, f"{path.name} 未声明 {code}_DISCLOSURE_SHEET_NAME"
    body = m.group("body")
    out: dict[str, str] = {}
    for variant in ("listed", "soe"):
        vm = re.search(rf"{variant}:\s*'([^']+)'", body)
        assert vm, f"{path.name} 缺 {variant} 项"
        out[variant] = vm.group(1)
    return out


def _read_xlsx_tabs(code: str) -> list[str]:
    hits = [
        p
        for p in TEMPLATE_DIR.glob(f"{code} *.xlsx")
        if not p.name.startswith("~$")
    ]
    assert len(hits) == 1, f"{code} 源 xlsx 命中 {len(hits)} 个：{[p.name for p in hits]}"
    wb = openpyxl.load_workbook(hits[0], read_only=True)
    try:
        return [s for s in wb.sheetnames if "附注" in s]
    finally:
        wb.close()


@pytest.mark.skipif(openpyxl is None, reason="openpyxl 未安装")
@pytest.mark.parametrize("code", CYCLES)
def test_sheet_name_matches_source_xlsx_tab(code: str) -> None:
    consts = _read_constants(code)
    tabs = _read_xlsx_tabs(code)
    assert tabs, f"{code} 源 xlsx 无披露 tab"

    listed_tab = next((t for t in tabs if "上市" in t), None)
    soe_tab = next((t for t in tabs if "国企" in t or "国有" in t), None)

    assert listed_tab is not None, f"{code} 源 xlsx 无上市披露 tab：{tabs}"
    assert soe_tab is not None, f"{code} 源 xlsx 无国企披露 tab：{tabs}"

    assert consts["listed"] == listed_tab, (
        f"{code}.listed 常量「{consts['listed']}」≠ 源 xlsx tab「{listed_tab}」"
        f"（括号半/全角必须原样，附注反向跳转按此精确匹配）"
    )
    assert consts["soe"] == soe_tab, (
        f"{code}.soe 常量「{consts['soe']}」≠ 源 xlsx tab「{soe_tab}」"
    )


@pytest.mark.skipif(openpyxl is None, reason="openpyxl 未安装")
def test_guard_is_not_vacuous() -> None:
    """反向自检：守卫必须真的能读到常量与 tab 名（防 glob / 正则失效后空跑）。"""
    consts = _read_constants("K2")
    assert consts["listed"].startswith("附注披露")
    tabs = _read_xlsx_tabs("K2")
    assert len(tabs) == 2


@pytest.mark.skipif(openpyxl is None, reason="openpyxl 未安装")
def test_detects_halfwidth_drift() -> None:
    """反向自检：把常量换成半角必须被判为漂移。"""
    tabs = _read_xlsx_tabs("K8")
    listed_tab = next(t for t in tabs if "上市" in t)
    drifted = listed_tab.replace("（", "(").replace("）", ")")
    assert drifted != listed_tab, "K8 源 tab 名应为全角，反向自检前提不成立"

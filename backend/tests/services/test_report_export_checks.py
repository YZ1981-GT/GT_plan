"""财务报表导出提醒（spec deliverable-trio-generation-accuracy AC-1.5 / AC-2.2）：不阻断导出，但不得静默。"""

from __future__ import annotations

from decimal import Decimal

from openpyxl import Workbook

from app.services.report_export_checks import (
    _check_balance,
    _check_equity_vs_bs,
    _check_impairment_source,
    _check_prior,
)
from app.services.report_template_binding import TemplateSlots


def _row(name, cur=None, prior=None):
    d = lambda v: None if v is None else Decimal(str(v))  # noqa: E731
    return {"row_code": "X-1", "row_name": name, "current_period_amount": d(cur), "prior_period_amount": d(prior)}


def test_balance_warns_with_difference_and_skips_all_zero():
    vals = {"资产总计": [(Decimal("100"), Decimal("90"))], "负债和所有者权益总计": [(Decimal("96"), Decimal("90"))]}
    out = _check_balance(vals)
    assert len(out) == 1 and "期末不平衡" in out[0] and "4.00" in out[0]
    assert _check_balance({"资产总计": [(Decimal(0), Decimal(0))], "负债和所有者权益总计": [(Decimal(0), Decimal(0))]}) == []


def test_impairment_all_zero_but_income_statement_has_losses():
    data = {
        "impairment_provision": [_row("一、坏账准备", 0, 0)],
        "income_statement": [_row("信用减值损失（损失以“－”号填列）", "13341759.48"), _row("营业收入", 9)],
    }
    out = _check_impairment_source(data)
    assert len(out) == 1 and "信用减值损失 13,341,759.48" in out[0]
    data["impairment_provision"] = [_row("一、坏账准备", 5, 1)]
    assert _check_impairment_source(data) == []


def test_prior_warnings_for_blank_opening_and_missing_income_prior():
    out = _check_prior({
        "balance_sheet": [_row("货币资金", 10, None)],
        "income_statement": [_row("营业收入", 9, None)],
    })
    assert any("期初余额为空" in w for w in out) and any("上期金额无来源" in w for w in out)
    assert _check_prior({"balance_sheet": [_row("货币资金", 10, 8)], "income_statement": [_row("营业收入", 9, 7)]}) == []


def test_equity_statement_year_end_vs_balance_sheet():
    wb = Workbook()
    ws = wb.active
    ws.title = "所有者权益变动表"
    ws["A6"], ws["C6"], ws["N6"] = "项目", "实收资本", "所有者权益合计"
    ws["A9"], ws["N9"] = "一、上年年末余额", 400
    ws["A13"], ws["N13"] = "二、本年年初余额", 400
    ws["A20"], ws["N20"] = "四、本年年末余额", 419
    slots = TemplateSlots()
    slots.sheet_family = {"所有者权益变动表": "EQ"}
    slots.row_labels = {("所有者权益变动表", 9): "一、上年年末余额", ("所有者权益变动表", 13): "二、本年年初余额",
                        ("所有者权益变动表", 20): "四、本年年末余额"}
    bs_vals = {"所有者权益合计": [(Decimal("398"), Decimal("400"))]}
    out = _check_equity_vs_bs(wb, slots, bs_vals)
    assert len(out) == 1 and "本年年末余额" in out[0] and "21.00" in out[0]
    ws["N20"] = 398
    assert _check_equity_vs_bs(wb, slots, bs_vals) == []

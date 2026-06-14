"""集成测试 — Word 导出 per-table + 章节级 guidance（note-per-table-guidance Phase 4）.

Spec:   .kiro/specs/note-per-table-guidance/ Phase 4 (Task 4.1 / 4.2)
Design: design.md §"Word 导出：note_word_exporter.py"
决策:   0.1 — guidance 进交付件，各表格前渲染 per-table guidance（区分样式）。

测试目标：
  1. 章节级 guidance_text 出现在生成 docx 中（此前完全不渲染）
  2. 表格级 _tables[0].guidance 出现在生成 docx 中
  3. guidance 段落使用区分样式（灰色 / 斜体 / 小字），与正式正文 run 区分
  4. 用 python-docx 读回生成文档断言（而非仅 XML 字符串）

策略：
- 沿用 test_note_multi_table_export.py 的 SimpleNamespace + Mock db 模式
- 用 python-docx Document(BytesIO) 读回，遍历 paragraph.runs 检查 font 属性
"""

from __future__ import annotations

from io import BytesIO
from types import SimpleNamespace
from unittest.mock import AsyncMock, MagicMock
from uuid import uuid4

import pytest
from docx import Document
from docx.shared import RGBColor

from app.services.note_word_exporter import (
    GUIDANCE_COLOR,
    GUIDANCE_FONT_SIZE,
    NoteWordExporter,
)

# ---------------------------------------------------------------------------
# 样本文字
# ---------------------------------------------------------------------------

SECTION_GUIDANCE = "（注：如有因抵押、质押或冻结等对使用有限制的款项，应单独披露。）"
TABLE0_GUIDANCE = "（提示：企业持有的货币资金应按存放地点分类列示。）"
BODY_TEXT = "本公司货币资金主要为银行存款及库存现金。"


# ---------------------------------------------------------------------------
# Helpers
# ---------------------------------------------------------------------------


def _make_note(
    section: str,
    title: str,
    *,
    table_data: dict | None = None,
    text_content: str | None = None,
    guidance_text: str | None = None,
):
    return SimpleNamespace(
        id=uuid4(),
        project_id=uuid4(),
        year=2025,
        note_section=section,
        section_title=title,
        account_name=None,
        content_type=None,
        table_data=table_data,
        text_content=text_content,
        guidance_text=guidance_text,
        source_template=None,
        status=None,
        sort_order=None,
        is_deleted=False,
        is_stale=False,
    )


def _make_db(notes: list) -> MagicMock:
    db = MagicMock()
    scalars = MagicMock()
    scalars.all = MagicMock(return_value=notes)
    result = MagicMock()
    result.scalars = MagicMock(return_value=scalars)
    db.execute = AsyncMock(return_value=result)
    return db


def _iter_runs(doc):
    """Yield (paragraph, run) over all body paragraphs."""
    for p in doc.paragraphs:
        for r in p.runs:
            yield p, r


def _find_run_with_text(doc, needle: str):
    """Return the first run whose text contains needle (or None)."""
    for _p, r in _iter_runs(doc):
        if needle in (r.text or ""):
            return r
    return None


def _is_guidance_styled(run) -> bool:
    """提示样式判定：灰色 + 斜体 + 小字（小于正文小四 12pt）。"""
    color_ok = run.font.color is not None and run.font.color.rgb == GUIDANCE_COLOR
    italic_ok = run.font.italic is True
    size_ok = run.font.size is not None and run.font.size == GUIDANCE_FONT_SIZE
    return color_ok and italic_ok and size_ok


# ---------------------------------------------------------------------------
# Task 4.2 — Word 导出 guidance 区分样式
# ---------------------------------------------------------------------------


@pytest.mark.asyncio
async def test_section_and_table_guidance_present_in_docx():
    """章节级 guidance_text + 表格级 _tables[0].guidance 均渲染进 docx."""
    note = _make_note(
        "八、1",
        "货币资金",
        text_content=BODY_TEXT,
        guidance_text=SECTION_GUIDANCE,
        table_data={
            "headers": ["项目", "期末余额", "期初余额"],
            "rows": [{"label": "库存现金", "values": [100.0, 90.0]}],
            "_tables": [
                {
                    "name": "货币资金",
                    "headers": ["项目", "期末余额", "期初余额"],
                    "rows": [{"label": "库存现金", "values": [100.0, 90.0]}],
                    "guidance": TABLE0_GUIDANCE,
                },
                {
                    "name": "受限制的货币资金明细",
                    "headers": ["项目", "期末余额", "期初余额"],
                    "rows": [{"label": "冻结存款", "values": [10.0, 5.0]}],
                },
            ],
        },
    )
    db = _make_db([note])
    exporter = NoteWordExporter(db)
    bio = await exporter.export(project_id=uuid4(), year=2025)

    doc = Document(BytesIO(bio.getvalue()))
    full_text = "\n".join(p.text for p in doc.paragraphs)

    assert SECTION_GUIDANCE in full_text, "章节级 guidance_text 未渲染进 docx"
    assert TABLE0_GUIDANCE in full_text, "表格级 _tables[0].guidance 未渲染进 docx"
    assert BODY_TEXT in full_text, "正文 text_content 应仍然渲染"


@pytest.mark.asyncio
async def test_guidance_runs_use_distinct_style():
    """guidance run 使用灰色 + 斜体 + 小字样式，与正文 run 区分."""
    note = _make_note(
        "八、1",
        "货币资金",
        text_content=BODY_TEXT,
        guidance_text=SECTION_GUIDANCE,
        table_data={
            "headers": ["项目", "期末余额", "期初余额"],
            "rows": [{"label": "库存现金", "values": [100.0, 90.0]}],
            "_tables": [
                {
                    "name": "货币资金",
                    "headers": ["项目", "期末余额", "期初余额"],
                    "rows": [{"label": "库存现金", "values": [100.0, 90.0]}],
                    "guidance": TABLE0_GUIDANCE,
                },
            ],
        },
    )
    db = _make_db([note])
    exporter = NoteWordExporter(db)
    bio = await exporter.export(project_id=uuid4(), year=2025)
    doc = Document(BytesIO(bio.getvalue()))

    section_run = _find_run_with_text(doc, SECTION_GUIDANCE)
    table_run = _find_run_with_text(doc, TABLE0_GUIDANCE)
    body_run = _find_run_with_text(doc, BODY_TEXT)

    assert section_run is not None and _is_guidance_styled(section_run), (
        "章节级 guidance 段落应为灰色/斜体/小字区分样式"
    )
    assert table_run is not None and _is_guidance_styled(table_run), (
        "表格级 guidance 段落应为灰色/斜体/小字区分样式"
    )

    # 正文 run 必须 NOT 是 guidance 样式（区分）
    assert body_run is not None
    assert not _is_guidance_styled(body_run), "正文 run 不应被误用 guidance 样式"


@pytest.mark.asyncio
async def test_no_guidance_no_extra_paragraph():
    """无 guidance（章节级与表格级均空）时不渲染任何提示段落（零回归）."""
    note = _make_note(
        "八、1",
        "货币资金",
        text_content=BODY_TEXT,
        guidance_text=None,
        table_data={
            "headers": ["项目", "期末余额"],
            "rows": [{"label": "库存现金", "values": [100.0]}],
        },
    )
    db = _make_db([note])
    exporter = NoteWordExporter(db)
    bio = await exporter.export(project_id=uuid4(), year=2025)
    doc = Document(BytesIO(bio.getvalue()))

    # 不应存在任何 guidance 样式 run
    for _p, r in _iter_runs(doc):
        assert not _is_guidance_styled(r), (
            f"无 guidance 时不应出现提示样式段落，但发现: {r.text!r}"
        )

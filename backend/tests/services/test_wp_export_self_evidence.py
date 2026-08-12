"""产物自证守卫 — workpaper-import-export-lifecycle-closure Wave 1 / Task 1

## 这个文件在守什么

用户反馈「导出的很多模板都是空的」。实证后确认根因是两个独立缺陷叠加，而**止损层**
是本文件的作业面：**空产物必须自证**，而不是伪装成「这份底稿内容本来就空」。

四档自证语义（`SelfEvidenceKind`）：

| kind                | 触发条件                                                        | 用户该做什么           |
|---------------------|-----------------------------------------------------------------|------------------------|
| `verdict`           | `resolve_wp_file` 返回 `verdict != 'file'`                      | 看 `VERDICT_LABELS` 原因 |
| `html_data_absent`  | `verdict == 'file'` 但 `html_data` 缺该 sheet 键                | 打开底稿保存一次       |
| `entry_not_in_xlsx` | `verdict == 'file'` + `html_data` 空 + `checklist_responses` 有行 | 录入在库里但没落 xlsx  |
| `None`              | `verdict == 'file'` + `html_data` 非空                          | 正常导出，**零噪声**   |

## 🔴 类 A / 类 B 分离（Wave 1 「先打红」设计）

- **类 A**（`TestExistingTruthSources`）：独立口径自检，**当前应全绿**。它验的是
  判据基础设施本身没写错（`VERDICT_LABELS` 四键齐备、替身自检有效）。若类 A 红，
  说明守卫自己有缺陷，不是被测实现的问题。
- **类 B**（其余 class）：被测实现，**Wave 1 当前应全红**，红消息带
  「尚未实现（Wave 2 Task 4）」。Task 4 交付 `self_evidence.py` 后应全绿。

## 🔴 为什么文案真源扫描要先剥注释

`build_self_evidence_banner` 的 docstring 里会出现示例文案（正是它要产出的中文），
若不剥注释，扫描「调用方是否硬写中文」时 docstring 会**冒充**成硬写而假红；反之
若剥得过头（把 `sa.text(\"\"\"...\"\"\"）` 一起剥掉）又会漏检。故本文件的
`_strip_py_comments` 只剥 `#` 行注释与模块/函数级 docstring，且配
`test_strip_comments_actually_works` 反向自检 —— memory 铁律「每写完守卫必做变异
检验，没打红=守卫有缺陷」。
"""

from __future__ import annotations

import ast
import importlib
import re
from pathlib import Path

import pytest

from app.services.wp_export.wp_file_resolver import (
    VERDICT_LABELS,
    WP_FILE_VERDICTS,
)

BACKEND_ROOT = Path(__file__).resolve().parents[2]

#: 尚未实现时的统一跳过/失败文案前缀（便于变异检验按名归因）
_NOT_YET = "尚未实现（Wave 2 Task 4）"

#: 自证共享件的模块路径（Task 4 交付物）
_SELF_EVIDENCE_MOD = "app.services.wp_export.self_evidence"

#: 必须接自证的调用方（Task 5/6 作业面）—— 文案真源扫描的扫描面
_CALLER_FILES: tuple[str, ...] = (
    "app/services/wp_download_service.py",
    "app/services/wp_xlsx_export_service.py",
    "app/services/wp_export/export_engine.py",
)


# ═══════════════════════════════════════════════════════════════════════════
# 辅助：源码剥注释（纯函数，便于反向自检）
# ═══════════════════════════════════════════════════════════════════════════


def _strip_py_comments(source: str) -> str:
    """剥掉 `#` 行注释与 docstring，保留其余字符串字面量。

    🔴 只剥 docstring（module/class/function 的首个裸字符串表达式），
    **不剥**普通字符串字面量 —— 否则 `sa.text(\"\"\"SELECT ...\"\"\")` 这类
    SQL 会被一起剥掉（memory 已记这个坑）。

    Returns:
        剥除后的源码。AST 解析失败时退回「只剥 # 行注释」，不抛。
    """
    # ── 先剥 docstring（用 AST 定位，按行号置空）────────────────────────
    lines = source.splitlines(keepends=True)
    try:
        tree = ast.parse(source)
    except SyntaxError:
        tree = None

    if tree is not None:
        doc_ranges: list[tuple[int, int]] = []
        for node in ast.walk(tree):
            if not isinstance(
                node, (ast.Module, ast.ClassDef, ast.FunctionDef, ast.AsyncFunctionDef)
            ):
                continue
            body = getattr(node, "body", None)
            if not body:
                continue
            first = body[0]
            if (
                isinstance(first, ast.Expr)
                and isinstance(first.value, ast.Constant)
                and isinstance(first.value.value, str)
                and first.lineno is not None
                and first.end_lineno is not None
            ):
                doc_ranges.append((first.lineno, first.end_lineno))
        for start, end in doc_ranges:
            for idx in range(start - 1, min(end, len(lines))):
                lines[idx] = "\n"
        source = "".join(lines)

    # ── 再剥 # 行注释（跳过字符串内的 #）────────────────────────────────
    out: list[str] = []
    for line in source.splitlines(keepends=True):
        in_s: str | None = None
        cut = None
        i = 0
        while i < len(line):
            ch = line[i]
            if in_s is None:
                if ch in ("'", '"'):
                    in_s = ch
                elif ch == "#":
                    cut = i
                    break
            else:
                if ch == "\\":
                    i += 2
                    continue
                if ch == in_s:
                    in_s = None
            i += 1
        out.append(line if cut is None else line[:cut].rstrip() + "\n")
    return "".join(out)


def _load_self_evidence():
    """import 自证共享件；未交付时 pytest.fail 带统一前缀（不 skip）。

    🔴 用 fail 而非 skip：skip 在 CI 里是绿的，会让「Wave 1 先打红」这个
    设计意图消失（memory 假绿三源之一 = 参数化集合为空导致整条 SKIP 空转）。
    """
    try:
        return importlib.import_module(_SELF_EVIDENCE_MOD)
    except ModuleNotFoundError:
        pytest.fail(f"{_NOT_YET}：缺 {_SELF_EVIDENCE_MOD}")


# ═══════════════════════════════════════════════════════════════════════════
# 类 A —— 独立口径自检（当前应全绿；红 = 守卫自身缺陷）
# ═══════════════════════════════════════════════════════════════════════════


class TestExistingTruthSources:
    """已有真源的结构自检 —— 不依赖 Task 4 交付物。"""

    def test_verdict_labels_cover_all_verdicts(self):
        """`VERDICT_LABELS` 必须覆盖 `WP_FILE_VERDICTS` 全部取值（R1.3 真源完整性）。"""
        assert set(WP_FILE_VERDICTS) == {
            "file",
            "template_fallback",
            "empty",
            "missing",
        }, f"verdict 取值域变了，自证四档需同步复核：{WP_FILE_VERDICTS}"
        missing = [v for v in WP_FILE_VERDICTS if v not in VERDICT_LABELS]
        assert not missing, f"VERDICT_LABELS 缺档: {missing}"

    def test_verdict_labels_are_distinct_chinese(self):
        """四档文案必须互不相同且是中文（否则用户看不出差别）。"""
        vals = [VERDICT_LABELS[v] for v in WP_FILE_VERDICTS]
        assert len(set(vals)) == len(vals), f"VERDICT_LABELS 文案有重复: {vals}"
        for v in vals:
            assert re.search(r"[\u4e00-\u9fff]", v), f"文案非中文: {v!r}"

    def test_strip_comments_actually_works(self):
        """反向自检：剥注释确实生效，且**不**剥普通字符串字面量。

        memory 铁律：源码守卫先 `stripComments()` + 配「剥注释确实生效」的自检。
        """
        src = (
            '"""模块 docstring 里的 本份为空白模板 应被剥掉。"""\n'
            "SQL = \"\"\"SELECT 1 FROM t\"\"\"\n"
            "X = 1  # 行注释里的 本份为空白模板 应被剥掉\n"
            'MSG = "普通字面量里的 本份为空白模板 必须保留"\n'
        )
        out = _strip_py_comments(src)
        assert "模块 docstring" not in out, "docstring 未被剥掉"
        assert "行注释里的" not in out, "# 行注释未被剥掉"
        assert "SELECT 1 FROM t" in out, "普通字符串被误剥（SQL 会被吃掉）"
        assert "普通字面量里的" in out, "普通字面量被误剥"
        # 剥后 docstring/注释里的自证文案应恰好消失两处、字面量里的保留一处
        assert out.count("本份为空白模板") == 1

    def test_caller_files_exist(self):
        """扫描面非空自检 —— 防「文件改名后扫描面为空导致整条空转」。"""
        for rel in _CALLER_FILES:
            assert (BACKEND_ROOT / rel).is_file(), f"扫描面缺失: {rel}"


# ═══════════════════════════════════════════════════════════════════════════
# 类 B —— 被测实现（Wave 1 应全红）
# ═══════════════════════════════════════════════════════════════════════════


class TestSelfEvidenceKinds:
    """R1.1 / R1.2 / R1.6 / R1.8：四档语义齐备且可判定。"""

    def test_module_exposes_four_state_api(self):
        mod = _load_self_evidence()
        for name in (
            "SelfEvidenceKind",
            "build_self_evidence_banner",
            "stamp_self_evidence",
            "needs_self_evidence",
        ):
            assert hasattr(mod, name), f"{_NOT_YET}：缺符号 {name}"

    def test_needs_self_evidence_returns_verdict_kind(self):
        """R1.1：`verdict != 'file'` ⇒ kind = 'verdict'。"""
        mod = _load_self_evidence()
        for verdict in ("template_fallback", "empty", "missing"):
            got = mod.needs_self_evidence(
                verdict=verdict, html_data={}, has_entry_rows=False
            )
            assert got == "verdict", f"verdict={verdict} 应判 'verdict'，实际 {got!r}"

    def test_needs_self_evidence_returns_html_absent(self):
        """R1.2：文件就绪但 html_data 缺键 ⇒ 'html_data_absent'。"""
        mod = _load_self_evidence()
        got = mod.needs_self_evidence(
            verdict="file", html_data={}, has_entry_rows=False
        )
        assert got == "html_data_absent", f"应判 'html_data_absent'，实际 {got!r}"

    def test_needs_self_evidence_returns_entry_not_in_xlsx(self):
        """R1.8：文件就绪 + html_data 空 + checklist 有行 ⇒ 独立成档。

        这是本 spec 最关键的一档：真实库 131 份底稿有 `checklist_responses` 行，
        而 `html_data` 非空的只有 73 份，两集合交集仅 5 份。绝大多数「有录入」
        的底稿走的正是这一档，必须与 `html_data_absent` 区分开 —— 前者要提示
        「录入在库里但未写进 xlsx」，后者要提示「打开底稿保存一次」。
        """
        mod = _load_self_evidence()
        got = mod.needs_self_evidence(
            verdict="file", html_data={}, has_entry_rows=True
        )
        assert got == "entry_not_in_xlsx", f"应判 'entry_not_in_xlsx'，实际 {got!r}"

    def test_needs_self_evidence_returns_none_on_normal_export(self):
        """R1.6：正常导出零噪声 —— 无条件替身自检（不依赖真实库）。"""
        mod = _load_self_evidence()
        got = mod.needs_self_evidence(
            verdict="file",
            html_data={"D2-1": {"rows": [{"a": 1}]}},
            has_entry_rows=True,
        )
        assert got is None, f"正常态必须返 None（禁给正常导出加噪声），实际 {got!r}"

    def test_entry_not_in_xlsx_and_html_absent_labels_differ(self):
        """R1.8：两档文案必须逐字不同（否则用户无法区分处置方式）。"""
        mod = _load_self_evidence()
        a = mod.build_self_evidence_banner(kind="html_data_absent")
        b = mod.build_self_evidence_banner(kind="entry_not_in_xlsx")
        assert a != b, "两档文案相同 ⇒ 区分档位失去意义"
        assert "录入内容未包含" in a or "录入内容未包含" in b, (
            "R1.2 要求产物首行写明「录入内容未包含」"
        )


class TestSelfEvidenceBanner:
    """R1.1 / R1.2：banner 文案内容判据。"""

    def test_verdict_banner_contains_blank_template_phrase_and_reason(self):
        """R1.1：含「本份为空白模板」+ `VERDICT_LABELS` 对应中文原因。"""
        mod = _load_self_evidence()
        for verdict in ("template_fallback", "empty", "missing"):
            banner = mod.build_self_evidence_banner(kind="verdict", verdict=verdict)
            assert "本份为空白模板" in banner, (
                f"verdict={verdict} 的 banner 缺「本份为空白模板」: {banner!r}"
            )
            assert VERDICT_LABELS[verdict] in banner, (
                f"verdict={verdict} 的 banner 未取 VERDICT_LABELS 原因: {banner!r}"
            )

    def test_html_absent_banner_has_actionable_hint(self):
        """R1.2：除原因外还须有「可操作提示」（不能只说坏消息）。"""
        mod = _load_self_evidence()
        banner = mod.build_self_evidence_banner(kind="html_data_absent")
        assert "录入内容未包含" in banner
        # 可操作提示的判据：出现动词性引导（打开/保存/重新生成/请）
        assert re.search(r"(请|打开|保存|重新生成)", banner), (
            f"缺可操作提示: {banner!r}"
        )


class TestSelfEvidenceStamping:
    """R1.7：自证行不占数据区首行语义。"""

    def test_stamp_puts_banner_row1_blank_row2_data_row3(self):
        """复用 `_build_failure_fallback_workbook` 范式：1 banner / 2 空 / 3 起数据。"""
        mod = _load_self_evidence()
        from openpyxl import Workbook

        wb = Workbook()
        ws = wb.active
        mod.stamp_self_evidence(ws, "测试 banner：本份为空白模板")
        assert ws.cell(row=1, column=1).value == "测试 banner：本份为空白模板"
        assert ws.cell(row=2, column=1).value in (None, ""), "第 2 行必须留空"
        assert getattr(mod, "SELF_EVIDENCE_DATA_START_ROW", None) == 3, (
            "数据起始行必须是常量 3（供调用方共用，禁各写一份行号）"
        )

    def test_banner_is_visually_marked(self):
        """banner 须加粗标红（与 `_build_failure_fallback_workbook` 一致）。"""
        mod = _load_self_evidence()
        from openpyxl import Workbook

        wb = Workbook()
        ws = wb.active
        mod.stamp_self_evidence(ws, "x")
        font = ws.cell(row=1, column=1).font
        assert font.bold, "banner 未加粗"
        assert (font.color is not None) and (
            str(font.color.rgb or "").upper().endswith("C00000")
        ), "banner 未标红（应与 _build_failure_fallback_workbook 同色 C00000）"


class TestSelfEvidenceSingleSource:
    """R1.3：自证文案单一真源，调用方不得硬写中文字面量。"""

    def test_kind_labels_is_the_only_extra_source(self):
        mod = _load_self_evidence()
        labels = getattr(mod, "_KIND_LABELS", None)
        assert isinstance(labels, dict), f"{_NOT_YET}：缺 _KIND_LABELS 单一真源"
        # 三档非 verdict 的 kind 至少覆盖两档（verdict 档文案来自 VERDICT_LABELS）
        for kind in ("html_data_absent", "entry_not_in_xlsx"):
            assert kind in labels, f"_KIND_LABELS 缺 {kind}"

    def test_callers_do_not_hardcode_self_evidence_text(self):
        """扫全部调用方源码，剥注释后不得出现硬写的自证中文字面量。

        判据取「自证专用短语」而非泛化中文 —— 泛化会把无关中文全打红。
        反向自检由 `test_strip_comments_actually_works` + 变异脚本承担。
        """
        mod = _load_self_evidence()
        banned = ["本份为空白模板", "录入内容未包含"]
        # 真源文件自己当然可以出现这些短语，扫描面已排除它
        self_evidence_rel = (
            Path(mod.__file__).resolve().relative_to(BACKEND_ROOT).as_posix()
        )

        offenders: list[str] = []
        for rel in _CALLER_FILES:
            if rel == self_evidence_rel:
                continue
            path = BACKEND_ROOT / rel
            src = _strip_py_comments(path.read_text(encoding="utf-8"))
            for phrase in banned:
                if phrase in src:
                    offenders.append(f"{rel} 硬写了 {phrase!r}")

        assert not offenders, (
            "自证文案必须只取 VERDICT_LABELS / _KIND_LABELS 单一真源，"
            f"以下调用方硬写了中文：{offenders}"
        )

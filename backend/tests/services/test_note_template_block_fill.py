"""附注模板模式按模板结构重建（spec deliverable-trio-generation-accuracy §2，note_template_block_fill）。

夹具按真实国企模板的块结构构造（2026-10-16 逐块核对 soe_standalone.docx）：``{{seq:八}}`` 独立成段 → 三级标题（显式
numPr 做章内编号重启）→ 四级子标题（首个显式重启）→ ``##STYLE_REF##`` → ``{{table:code:N}}`` + 样例表 →
``{{section:code}}`` → 编制指引 → 块结束标记之后的参考表。
"""
from __future__ import annotations

import copy
from types import SimpleNamespace

import pytest
from docx import Document
from docx.oxml import OxmlElement
from docx.oxml.ns import qn

from app.services import note_template_block_fill as tbf
from app.services.note_deliverable_filter import FilterStats, make_block_filter
from app.services.word_doc_utils import scan_section_blocks

W = "{http://schemas.openxmlformats.org/wordprocessingml/2006/main}"


def _num(p, num_id: str, ilvl: str = "2") -> None:
    ppr = p._p.get_or_add_pPr()
    np_ = OxmlElement("w:numPr")
    il = OxmlElement("w:ilvl")
    il.set(qn("w:val"), ilvl)
    ni = OxmlElement("w:numId")
    ni.set(qn("w:val"), num_id)
    np_.append(il)
    np_.append(ni)
    ppr.append(np_)


def _numid(el) -> str | None:
    n = el.find(f"{W}pPr/{W}numPr/{W}numId")
    return n.get(f"{W}val") if n is not None else None


def _sample_table(doc, header: list[str], rows: int = 2):
    """模板样例表：通栏（tblW 5000 pct）、居中 —— 同真实模板。"""
    t = doc.add_table(rows=rows, cols=len(header))
    for i, h in enumerate(header):
        t.cell(0, i).text = h
    tblw = t._tbl.tblPr.find(qn("w:tblW"))
    tblw.set(qn("w:w"), "5000")
    tblw.set(qn("w:type"), "pct")
    return t


def _texts(doc) -> list[str]:
    return ["".join(t.text or "" for t in p.iter(f"{W}t")) for p in doc.element.body.iter(f"{W}p")]


def _style_numbering(doc, style_name: str, num_id: str, ilvl: str) -> None:
    """样式带自动编号（真实模板的标题样式都带 numPr）。"""
    st = doc.styles[style_name].element
    ppr = st.find(f"{W}pPr")
    if ppr is None:
        ppr = OxmlElement("w:pPr")
        st.append(ppr)
    np_ = OxmlElement("w:numPr")
    for tag, val in (("w:ilvl", ilvl), ("w:numId", num_id)):
        el = OxmlElement(tag)
        el.set(qn("w:val"), val)
        np_.append(el)
    ppr.insert(0, np_)


def _block_doc():
    """一章（一级标题）两块：八、1 货币资金（主表 + 块尾参考表）、八、4 应收票据（子标题 + 多表位）。"""
    doc = Document()
    _style_numbering(doc, "Heading 3", "3", "2")
    _style_numbering(doc, "Heading 4", "4", "3")
    doc.add_paragraph("财务报表主要项目注释", style="Heading 1")
    doc.add_paragraph("##SECTION:八、1##")
    doc.add_paragraph("{{seq:八}}")
    t1 = doc.add_paragraph("货币资金", style="Heading 3")
    _num(t1, "24")  # 章内节编号重启
    doc.add_paragraph("##STYLE_REF:table:八、1:0##")
    doc.add_paragraph("{{table:八、1:0}}")
    _sample_table(doc, ["项  目", "期末余额", "期初余额"])
    doc.add_paragraph("{{section:八、1}}")
    doc.add_paragraph("（注：如有因抵押、质押或冻结等对使用有限制的款项应单独说明。）")
    doc.add_paragraph("受限制的货币资金明细")
    doc.add_paragraph("##/SECTION:八、1##")
    _sample_table(doc, ["项  目", "期末余额", "期初余额"])  # 块结束标记之后的参考表
    doc.add_paragraph("##SECTION:八、4##")
    doc.add_paragraph("{{seq:八}}")
    doc.add_paragraph("应收票据", style="Heading 3")
    s1 = doc.add_paragraph("应收票据分类", style="Heading 4")
    _num(s1, "25", "3")  # 块内子标题编号重启
    doc.add_paragraph("{{table:八、4:0}}")
    _sample_table(doc, ["票据种类", "期末数", "期初数"])
    doc.add_paragraph("坏账准备计提情况", style="Heading 4")
    doc.add_paragraph("{{section:八、4}}")
    doc.add_paragraph("按坏账准备计提方法分类披露应收票据")
    doc.add_paragraph("{{table:八、4:1}}")
    _sample_table(doc, ["类  别", "期末数"])
    doc.add_paragraph("期末已质押的应收票据", style="Heading 4")
    doc.add_paragraph("{{table:八、4:2}}")
    _sample_table(doc, ["票据种类", "期末已质押金额"])
    doc.add_paragraph("##/SECTION:八、4##")
    return doc


class _Exporter:
    """``rebuild_block`` 只用到导出器的表投影与 ``_render_table``：用真实导出器（不连库）。"""

    def __new__(cls):
        from app.services.note_word_exporter import NoteWordExporter

        return NoteWordExporter(db=None)


def _table(name: str, label: str, value: float, headers=("项目", "期末余额")) -> dict:
    return {"name": name, "headers": list(headers), "rows": [{"label": label, "values": [value]},
                                                              {"label": "合计", "values": [value]}]}


def _note(code: str, title: str, text: str = "", tables: list[dict] | None = None):
    return SimpleNamespace(note_section=code, section_title=title, text_content=text,
                           table_data={"_tables": tables} if tables else None, source_template=None,
                           is_deleted=False, status=None, is_empty=False)


def _blocks(doc):
    tbf.absorb_trailing_content(doc)
    return {b.section_code: b for b in scan_section_blocks(doc)}


def test_layout_reads_template_block_structure():
    doc = _block_doc()
    blocks = _blocks(doc)
    names = tbf.style_map(doc)
    lay = tbf.block_layout(blocks["八、4"], names)
    assert tbf.paragraph_text(lay.title) == "应收票据"
    assert [tbf.paragraph_text(s) for s in lay.subheadings] == ["应收票据分类", "坏账准备计提情况", "期末已质押的应收票据"]
    assert [(s.index, tbf.paragraph_text(s.caption), tbf.paragraph_text(s.subheading)) for s in lay.slots] == [
        (0, "", "应收票据分类"),
        (1, "按坏账准备计提方法分类披露应收票据", "坏账准备计提情况"),
        (2, "", "期末已质押的应收票据"),
    ]
    assert tbf.paragraph_text(lay.section_sub) == "坏账准备计提情况"
    # 块尾参考表收进了八、1：它成了第二个表位，表题是「受限制的货币资金明细」（指引段「（注：…）」不当表题）
    lay1 = tbf.block_layout(blocks["八、1"], names)
    assert [tbf.paragraph_text(s.caption) for s in lay1.slots] == ["", "受限制的货币资金明细"]


def _rebuild(doc, code: str, note, stats: FilterStats | None = None):
    blocks = _blocks(doc)
    return tbf.rebuild_block(doc, blocks[code], note, exporter=_Exporter(),
                             text_filter=make_block_filter(note, stats or FilterStats()))


def _block_texts(doc, code: str) -> list[str]:
    b = next(x for x in scan_section_blocks(doc) if x.section_code == code)
    out = []
    for el in b.elements[1:-1]:
        out.extend("".join(t.text or "" for t in p.iter(f"{W}t")) for p in ([el] if el.tag == f"{W}p" else el.iter(f"{W}p")))
    return [t for t in out if t]


def test_rebuild_keeps_title_tables_text_and_drops_template_content():
    doc = _block_doc()
    note = _note("八、1", "货币资金", text="八、1 货币资金\n本公司货币资金期末余额中无受限资金。\n审计师提示：请补充明细。",
                 tables=[_table("货币资金", "银行存款", 100.0),
                         _table("受限制的货币资金明细", "保证金", 50.0)])
    st = FilterStats()
    rebuilt = _rebuild(doc, "八、1", note, st)
    texts = _block_texts(doc, "八、1")
    # 标题 + 主表 + 正文 + 块尾表位（模板表题 + 库内表）；{{seq}}、标记、指引、样例表都不在
    assert texts[0] == "货币资金"
    assert "银行存款" in texts and "保证金" in texts and "受限制的货币资金明细" in texts
    assert "本公司货币资金期末余额中无受限资金。" in texts
    assert not [t for t in texts if t.startswith(("{{", "##", "（注")) or "审计师" in t or t == "八、1 货币资金"]
    assert texts.index("银行存款") < texts.index("本公司货币资金期末余额中无受限资金。") < texts.index("保证金")
    assert rebuilt.stats.tables_rendered == 2 and rebuilt.stats.extra_tables == 0
    assert st.dropped_lines == 2 and st.hint_tail_cut  # 重复标题行 + 提示段（及其后）
    b = next(x for x in scan_section_blocks(doc) if x.section_code == "八、1")
    tables = [el for el in b.elements if el.tag == f"{W}tbl"]
    assert len(tables) == 2  # 模板两张样例表已换成库内表


def test_unmatched_slots_drop_their_subheadings_and_restart_moves_to_next_kept():
    doc = _block_doc()
    # 只有第二个表位（按表名对上）有数：「应收票据分类」子标题（带编号重启）随表位删除，重启转给「坏账准备计提情况」
    note = _note("八、4", "应收票据", tables=[_table("按坏账准备计提方法分类披露应收票据", "银行承兑汇票", 10.0)])
    _rebuild(doc, "八、4", note)
    b = next(x for x in scan_section_blocks(doc) if x.section_code == "八、4")
    paras = [el for el in b.elements if el.tag == f"{W}p"]
    by_text = {"".join(t.text or "" for t in p.iter(f"{W}t")): p for p in paras}
    assert "应收票据分类" not in by_text and "期末已质押的应收票据" not in by_text
    assert _numid(by_text["坏账准备计提情况"]) == "25"
    assert "按坏账准备计提方法分类披露应收票据" in by_text  # 模板表题保留


def test_extra_and_continuation_tables_get_captions_from_table_names():
    doc = _block_doc()
    note = _note("八、4", "应收票据", tables=[
        _table("应收票据分类", "银行承兑汇票", 10.0),
        _table("应收票据分类（续）", "商业承兑汇票", 5.0),  # 模板无续表表位 → 紧跟主表、按表名补表题
        _table("模板里没有的表", "其他", 1.0),  # 无表位 → 追加在最后一个保留表之后
    ])
    rebuilt = _rebuild(doc, "八、4", note)
    texts = _block_texts(doc, "八、4")
    assert texts.index("应收票据分类") < texts.index("银行承兑汇票") < texts.index("应收票据分类（续）") \
        < texts.index("商业承兑汇票") < texts.index("模板里没有的表") < texts.index("其他")
    assert rebuilt.stats.extra_tables == 1


def test_empty_tables_are_not_rendered_and_text_goes_after_title_without_anchor():
    doc = Document()
    doc.add_paragraph("##SECTION:八、2##")
    doc.add_paragraph("交易性金融资产", style="Heading 3")
    doc.add_paragraph("{{table:八、2:0}}")
    _sample_table(doc, ["项  目", "期末公允价值"])
    doc.add_paragraph("##/SECTION:八、2##")
    empty = {"name": "交易性金融资产", "headers": ["项目", "期末"], "rows": [{"label": "债务工具投资", "values": [None]}]}
    note = _note("八、2", "交易性金融资产", text="本公司期末持有的交易性金融资产为理财产品。", tables=[empty])
    _rebuild(doc, "八、2", note)
    assert _block_texts(doc, "八、2") == ["交易性金融资产", "本公司期末持有的交易性金融资产为理财产品。"]
    b = next(x for x in scan_section_blocks(doc) if x.section_code == "八、2")
    assert not [el for el in b.elements if el.tag == f"{W}tbl"]


def test_literal_subheading_numbers_are_renumbered_in_kept_order():
    doc = Document()
    doc.add_paragraph("##SECTION:八、5##")
    doc.add_paragraph("应收账款", style="Heading 3")
    for i, name in enumerate(["按账龄披露", "按坏账准备计提方法分类", "本期计提坏账准备情况"], 1):
        p = doc.add_paragraph(f"（{i}）{name}", style="Heading 4")
        _num(p, "0", "3")  # 字面编号：取消了自动编号
        doc.add_paragraph(f"{{{{table:八、5:{i - 1}}}}}")
        _sample_table(doc, ["项目", "期末"])
    doc.add_paragraph("##/SECTION:八、5##")
    note = _note("八、5", "应收账款", tables=[_table("按坏账准备计提方法分类", "组合", 1.0),
                                          _table("本期计提坏账准备情况", "坏账准备", 2.0)])
    _rebuild(doc, "八、5", note)
    texts = _block_texts(doc, "八、5")
    assert "（1）按坏账准备计提方法分类" in texts and "（2）本期计提坏账准备情况" in texts
    assert not [t for t in texts if "按账龄披露" in t]


def _literal_block(code: str, subs: list[str], *, auto_first: bool = False):
    """块内字面编号子标题（取消自动编号）各带一个表位；``auto_first`` = 第一个子标题靠样式自动编号。"""
    doc = Document()
    _style_numbering(doc, "Heading 4", "4", "3")
    doc.add_paragraph(f"##SECTION:{code}##")
    doc.add_paragraph("标题", style="Heading 3")
    for i, text in enumerate(subs):
        p = doc.add_paragraph(text, style="Heading 4")
        if not (auto_first and i == 0):
            _num(p, "0", "3")
        doc.add_paragraph(f"{{{{table:{code}:{i}}}}}")
        _sample_table(doc, ["项目", "期末"])
    doc.add_paragraph(f"##/SECTION:{code}##")
    return doc


def test_continuation_caption_strips_guidance_and_follows_its_subheading_number():
    """续表表题按表名补：剔除编制指引括注；所属子标题「（2）」重排成「（1）」时表题编号跟着改（spec
    note-pull-header-formula-closure R5.4）。"""
    doc = _literal_block("八、5", ["（1）按账龄披露", "（2）按坏账准备计提方法分类", "（3）本期计提坏账准备情况"])
    note = _note("八、5", "应收账款", tables=[
        _table("（2）按坏账准备计提方法分类", "组合", 1.0),
        _table("（2）按坏账准备计提方法分类（续：期初数）（注意：是报表项目）", "组合", 3.0),
        _table("（3）本期计提坏账准备情况", "坏账准备", 2.0),
    ])
    _rebuild(doc, "八、5", note)
    texts = _block_texts(doc, "八、5")
    assert "（1）按坏账准备计提方法分类" in texts and "（2）本期计提坏账准备情况" in texts
    assert "（1）按坏账准备计提方法分类（续：期初数）" in texts, texts
    assert not [t for t in texts if "注意" in t or "按账龄披露" in t]


def test_literal_numbers_not_consecutive_from_one_in_template_are_kept():
    """上市 五、71「（7）供应商融资安排」接在自动编号子标题后：模板里该组不是从 1 连续 ⇒ 不重排（原先成了「（1）」）。"""
    doc = _literal_block("五、71", ["现金流量表补充资料", "（7）供应商融资安排"], auto_first=True)
    note = _note("五、71", "现金流量表补充资料", tables=[_table("供应商融资安排", "应付账款", 1.0)])
    _rebuild(doc, "五、71", note)
    assert "（7）供应商融资安排" in _block_texts(doc, "五、71")


def test_caption_text_strips_only_guidance():
    assert tbf.caption_text("应收关联方款项（注意：是报表项目）") == "应收关联方款项"
    assert tbf.caption_text("资产负债表日后利润分配情况（只写现金分配，若仅有转增，则不写）") == "资产负债表日后利润分配情况"
    assert tbf.caption_text("营业外收入（损失以“—”号填列）") == "营业外收入（损失以“—”号填列）"
    assert tbf.caption_text("（参考格式同上）") == "（参考格式同上）"  # 只剩指引时原样（不写空表题）


def test_template_table_look_header_repeat_min_height_and_total_bold():
    doc = _block_doc()
    note = _note("八、1", "货币资金", tables=[_table("货币资金", "银行存款", 100.0)])
    _rebuild(doc, "八、1", note)
    b = next(x for x in scan_section_blocks(doc) if x.section_code == "八、1")
    tbl = next(el for el in b.elements if el.tag == f"{W}tbl")
    rows = tbl.findall(f"{W}tr")
    assert rows[0].find(f"{W}trPr/{W}tblHeader") is not None  # 表头跨页重复
    assert all(r.find(f"{W}trPr/{W}trHeight").get(f"{W}hRule") == "atLeast" for r in rows)
    assert tbl.find(f"{W}tblPr/{W}tblW").get(f"{W}type") == "pct"
    total = rows[-1]
    assert all(r.find(f"{W}rPr/{W}b") is not None for r in total.iter(f"{W}r"))  # 合计行加粗
    assert not list(tbl.iter(f"{W}sz"))  # 3 列表：去掉固定字号，继承模板正文字号（样例表未写字号）


def _amount_table(cols: int):
    doc = Document()
    t = doc.add_table(rows=2, cols=cols)
    for c in range(cols):
        t.cell(0, c).text = f"列{c}"
        t.cell(1, c).text = "79,402,464.61" if c else "其中：固定资产"
    return t._tbl


def _grid(tbl) -> list[int]:
    return [int(g.get(f"{W}w")) for g in tbl.findall(f"{W}tblGrid/{W}gridCol")]


def test_amount_columns_widen_from_label_column():
    tbl = _amount_table(5)
    tbf.apply_template_table_look(tbl, None, header_rows=1)
    grid = _grid(tbl)
    # 5 列：继承模板正文字号；等分时金额列 1400 twip 放不下「79,402,464.61」（五号约 1486），从首列匀出
    assert not list(tbl.iter(f"{W}sz"))
    assert min(grid[1:]) > 1400 and grid[0] < 2800 and abs(sum(grid) - 8400) <= 5


def test_very_wide_table_falls_back_to_8pt_with_label_floor():
    tbl = _amount_table(7)
    tbf.apply_template_table_look(tbl, None, header_rows=1)
    grid = _grid(tbl)
    assert {s.get(f"{W}val") for s in tbl.iter(f"{W}sz")} == {"16"}  # 9 磅仍放不下 → 8 磅
    assert grid[0] == 1300 and sum(grid) <= 8400


def test_title_restart_moves_to_first_kept_title_in_chapter():
    doc = _block_doc()
    names, styles = tbf.style_map(doc), tbf.styles_by_id(doc)
    blocks = _blocks(doc)
    layouts = {c: tbf.block_layout(b, names) for c, b in blocks.items()}
    order = tbf.capture_title_order(doc, layouts.values(), names)
    tbf.keep_only_h1(blocks["八、1"], names)  # 「货币资金」（带章内重启 numId=24）整块删
    tbf.rebuild_block(doc, blocks["八、4"], _note("八、4", "应收票据", text="本公司应收票据均为银行承兑汇票。"),
                      exporter=_Exporter(), names=names, styles=styles, layout=layouts["八、4"])
    assert tbf.transfer_title_restarts(order, styles) == 1
    title = next(p for p in doc.element.body.iter(f"{W}p") if "".join(t.text or "" for t in p.iter(f"{W}t")) == "应收票据")
    assert _numid(title) == "24"
    assert "财务报表主要项目注释" in _texts(doc)  # 一级标题不随块删除


def test_keep_only_h1_keeps_heading_and_section_break():
    doc = Document()
    doc.add_paragraph("##SECTION:一、1##")
    doc.add_paragraph("公司基本情况", style="Heading 1")
    p = doc.add_paragraph("XX有限公司（以下简称本公司）……")
    sect = OxmlElement("w:sectPr")
    p._p.get_or_add_pPr().append(sect)
    doc.add_paragraph("##/SECTION:一、1##")
    names = tbf.style_map(doc)
    tbf.keep_only_h1(scan_section_blocks(doc)[0], names)
    assert [t for t in _texts(doc) if t] == ["公司基本情况"]
    assert doc.element.body.find(f".//{W}pPr/{W}sectPr") is not None  # 分节保留（只清文字）


def test_loose_regions_guidance_and_boilerplate():
    from app.services.note_deliverable_outline import parse_note_template_outline

    doc = Document()
    doc.add_paragraph("财务报表项目附注", style="Heading 1")
    doc.add_paragraph("【企业应当在附注中披露…】")  # 章首、块外的编制指引
    doc.add_paragraph("##SECTION:五、1##")
    doc.add_paragraph("货币资金", style="Heading 3")
    doc.add_paragraph("{{section:五、1}}")
    doc.add_paragraph("##/SECTION:五、1##")
    doc.add_paragraph("公允价值变动收益", style="Heading 3")  # 块外同级标题
    doc.add_paragraph("{{section:五、loose}}")
    doc.add_paragraph("财务报表之批准", style="Heading 1")
    doc.add_paragraph("本财务报表业经本公司董事会于202X年×月×日批准报出。")
    doc.add_paragraph("××××有限公司")
    names = tbf.style_map(doc)
    outline = parse_note_template_outline(doc)
    loose = tbf.wrap_loose_regions(doc, outline)
    assert [r.title for r in loose.values()] == ["公允价值变动收益"]
    tbf.absorb_trailing_content(doc)
    assert tbf.strip_loose_guidance(doc, names) == 1
    pending = tbf.fill_boilerplate(doc, names, "测试科技有限公司")
    texts = [t for t in _texts(doc) if t]
    assert "【企业应当在附注中披露…】" not in texts
    assert "测试科技有限公司" in texts and pending == ["财务报表之批准"]  # 日期仍待填写
    assert [b.section_code for b in scan_section_blocks(doc)] == ["五、1", next(iter(loose))]


def test_header_placeholders_replaced_across_split_runs():
    doc = Document()
    p = doc.sections[0].header.paragraphs[0]
    for part in ("{{company_", "full_name}}", " 财务报表附注 ", "{{audit_year}}年度"):
        p.add_run(part)
    assert tbf.replace_doc_placeholders(doc, {"company_full_name": "测试科技有限公司", "audit_year": "2025"}) >= 1
    assert p.text == "测试科技有限公司 财务报表附注 2025年度"


def test_cross_refs_rewritten_to_delivered_numbers():
    doc = Document()
    a = doc.add_paragraph("详见附注八、10 存货及附注八、99。")
    b = doc.add_paragraph()
    for part in ("见附注", "八、", "1", "（1）"):
        b.add_run(part)
    changed, dangling = tbf.rewrite_cross_refs([a._p, b._p], {"八、10": "七、9", "八、1": "七、1"})
    assert changed == 2 and dangling == ["八、99"]
    assert a.text == "详见附注七、9 存货及附注八、99。" and b.text == "见附注七、1（1）"


def test_emptied_landscape_section_is_collapsed():
    doc = Document()
    doc.add_paragraph("正文一")
    p1 = doc.add_paragraph()  # 纵向节结束
    p1._p.get_or_add_pPr().append(copy.deepcopy(doc.sections[0]._sectPr))
    p2 = doc.add_paragraph()  # 横向节结束（节内容已被删光）
    land = copy.deepcopy(doc.sections[0]._sectPr)
    land.find(f"{W}pgSz").set(f"{W}orient", "landscape")
    p2._p.get_or_add_pPr().append(land)
    doc.add_paragraph("正文二")
    assert tbf.collapse_empty_sections(doc) == 2  # 空的横向节并掉，前后同版式的纵向节随之合并
    assert doc.element.body.findall(f".//{W}pPr/{W}sectPr") == []
    assert [t for t in _texts(doc) if t] == ["正文一", "正文二"]


def test_cell_only_paragraph_is_cleared_not_deleted_by_marker_cleanup():
    """收尾清理遇到单元格里唯一的段落只清空（``w:tc`` 至少要一个段落，否则 Word 报内容损坏）。"""
    from app.services.word_doc_utils import remove_section_markers

    doc = Document()
    doc.add_table(rows=1, cols=1).cell(0, 0).paragraphs[0].add_run("##STYLE_REF:table:八、1:1##")
    remove_section_markers(doc)
    cells = list(doc.element.body.iter(f"{W}tc"))
    assert cells and all(c.find(f"{W}p") is not None for c in cells)
    assert not [t for t in _texts(doc) if t.startswith("##")]


# ── 真实模板：附注交付件编号 = 报表附注列编号（同一保留谓词 + 同一轮廓） ──────────────────────


class _Nested:
    async def __aenter__(self):
        return self

    async def __aexit__(self, *exc):
        return False


class _Result:
    def __init__(self, *, rows=None, scalars=None, first=None):
        self._rows, self._scalars, self._first = rows or [], scalars or [], first

    def scalars(self):
        return SimpleNamespace(all=lambda: list(self._scalars))

    def all(self):
        return list(self._rows)

    def first(self):
        return self._first


class _FakeDB:
    """按查询实体回数：附注 / 报表行 / 项目名（其余空）。"""

    def __init__(self, notes, fs_rows):
        self.notes, self.fs_rows = notes, fs_rows

    def begin_nested(self):
        return _Nested()

    async def execute(self, stmt):
        from app.models.core import Project
        from app.models.report_models import DisclosureNote, FinancialReport

        ents = {d.get("entity") for d in getattr(stmt, "column_descriptions", [])}
        if DisclosureNote in ents:
            return _Result(scalars=self.notes)
        if FinancialReport in ents:
            return _Result(rows=self.fs_rows)
        if Project in ents:
            return _Result(first=("测试科技有限公司", "测试科技_2025"))
        return _Result()


def _real_notes():
    cash = _note("八、1", "货币资金", text="货币资金的受限情况详见附注八、3。",
                 tables=[_table("货币资金", "银行存款", 100.0)])
    trading = _note("八、2", "交易性金融资产", text="本公司交易性金融资产为银行理财产品。")  # 报表余额 0、只有文字
    deriv = _note("八、3", "衍生金融资产", tables=[_table("衍生金融资产", "远期外汇合约", 5.0)])
    return [cash, trading, deriv]


_FS_ROWS = [("货币资金", 100.0, 80.0), ("交易性金融资产", 0.0, 0.0), ("衍生金融资产", 5.0, 0.0)]


@pytest.mark.asyncio
async def test_delivered_note_numbers_match_fs_note_column():
    from app.services.note_deliverable_outline import build_fs_note_ref_plan, load_note_template_outline
    from app.services.note_word_exporter import NoteWordExporter

    if load_note_template_outline("soe", "standalone") is None:
        pytest.skip("真实附注模板缺失")
    pid = SimpleNamespace()
    db = _FakeDB(_real_notes(), _FS_ROWS)
    buf, meta = await NoteWordExporter(db).export_with_meta(
        pid, 2025, template_type="soe", report_scope="standalone", mode="template")
    plan = await build_fs_note_ref_plan(_FakeDB(_real_notes(), _FS_ROWS), pid, 2025,
                                        template_type="soe", report_scope="standalone")
    # 单体国企模板没有合并专用章：「财务报表主要项目注释」显示为第七章；八、2 报表零余额只有文字 → 不出具，八、3 顺延为七、2
    assert plan.ref_for("货币资金") == "七、1"
    assert plan.ref_for("衍生金融资产") == "七、2"
    assert plan.ref_for("交易性金融资产") is None
    texts = _texts(Document(buf))
    assert "货币资金的受限情况详见附注七、2。" in texts  # 正文交叉引用与报表附注列同一编号
    assert not [t for t in texts if "交易性金融资产为银行理财产品" in t]
    assert set(meta.kept_codes) >= {"八、1", "八、3"} and "八、2" not in meta.kept_codes
    assert any("八、2 交易性金融资产" in w and "余额为 0" in w for w in meta.delivery_warnings)
    assert "测试科技有限公司" in Document(buf).sections[0].header.paragraphs[0].text


def test_guidance_bracket_before_table_is_not_a_caption():
    """「【披露要求：…4.应披露…累计金额。】」紧贴表位：不当表题（首版把指引尾巴当表题输出）。"""
    doc = Document()
    doc.add_paragraph("##SECTION:八、12##")
    doc.add_paragraph("持有待售资产", style="Heading 3")
    doc.add_paragraph("【披露要求：")
    doc.add_paragraph("1.应披露持有待售的非流动资产或处置组的出售费用；")
    doc.add_paragraph("4.应披露与持有待售的非流动资产或处置组有关的其他综合收益累计金额。】")
    doc.add_paragraph("{{table:八、12:0}}")
    _sample_table(doc, ["项目", "期末数", "期初数"])
    doc.add_paragraph("（注：如有因抵押、质押对使用有限制的资产，应单独说明。）")  # 括号外的指引段同样不当表题
    doc.add_paragraph("{{table:八、12:1}}")
    _sample_table(doc, ["项目", "期末数"])
    doc.add_paragraph("##/SECTION:八、12##")
    names = tbf.style_map(doc)
    lay = tbf.block_layout(scan_section_blocks(doc)[0], names)
    assert [s.caption for s in lay.slots] == [None, None]
    _rebuild(doc, "八、12", _note("八、12", "持有待售资产", tables=[_table("持有待售资产", "固定资产", 79402464.61)]))
    texts = _block_texts(doc, "八、12")
    assert texts[0] == "持有待售资产" and not [t for t in texts if "应披露" in t or "【" in t]


@pytest.mark.asyncio
async def test_chapter_level_note_fills_empty_chapter_only_when_titles_match():
    """章级附注（编号只有章号）：该章模板块都没有内容时放到章标题后；章名对不上（上市 / 国企章号含义不同）不放。"""
    from app.services.note_deliverable_outline import load_note_template_outline
    from app.services.note_word_exporter import NoteWordExporter

    if load_note_template_outline("soe", "standalone") is None:
        pytest.skip("真实附注模板缺失")
    notes = [
        _note("九", "或有事项", text="截至 2025 年 12 月 31 日，本公司无需要披露的重大或有事项。"),
        _note("十", "其他重要事项", text="本公司无其他重要事项。"),  # 国企单体第九章是「资产负债表日后事项」：章名不符
    ]
    buf, meta = await NoteWordExporter(_FakeDB(notes, [])).export_with_meta(
        SimpleNamespace(), 2025, template_type="soe", report_scope="standalone", mode="template")
    texts = [t for t in _texts(Document(buf)) if t]
    i = texts.index("或有事项")
    assert texts[i + 1] == "截至 2025 年 12 月 31 日，本公司无需要披露的重大或有事项。"
    assert "本公司无其他重要事项。" not in texts
    assert "九" in meta.kept_codes and "十" not in meta.kept_codes
    # 章级附注单列一条（交付中心勾选树里叫「本章整体说明」），不混进「模板没有对应位置」那条
    assert any("本章整体说明" in w and "十 其他重要事项" in w for w in meta.delivery_warnings), meta.delivery_warnings
    assert not [w for w in meta.delivery_warnings if "没有对应位置" in w and "十 其他重要事项" in w]


def test_note_display_label_does_not_repeat_name_coded_titles():
    """导出提醒的章节称呼：会计政策等按名称编码的章节（「四、存货」）不重复成「四、存货 存货」（Playwright 实测提醒原文）。"""
    from app.services.note_template_export_fill import note_display_label

    assert note_display_label("四、存货", "存货") == "四、存货"
    assert note_display_label("四、优先股、永续债等其他金融工具", "优先股、永续债等其他金融工具") == \
        "四、优先股、永续债等其他金融工具"
    assert note_display_label("八、2", "交易性金融资产") == "八、2 交易性金融资产"
    assert note_display_label("十", "其他重要事项") == "十 其他重要事项"
    assert note_display_label("八、9", None) == "八、9"
    assert note_display_label(" 八、9 ", "  ") == "八、9"


@pytest.mark.asyncio
async def test_name_coded_section_listed_once_in_delivery_warnings():
    """端到端：按名称编码、已勾选但没有可出具内容的政策章节，在导出提醒里只出现编号一次。"""
    from app.services.note_deliverable_outline import load_note_template_outline
    from app.services.note_word_exporter import NoteWordExporter

    if load_note_template_outline("soe", "standalone") is None:
        pytest.skip("真实附注模板缺失")
    notes = [
        _note("四、存货", "存货", text="存货\n【编制说明：请按企业实际填写存货计价方法。】"),
        _note("九", "或有事项", text="截至 2025 年 12 月 31 日，本公司无需要披露的重大或有事项。"),
    ]
    _buf, meta = await NoteWordExporter(_FakeDB(notes, [])).export_with_meta(
        SimpleNamespace(), 2025, template_type="soe", report_scope="standalone", mode="template",
        sections=["四、存货", "九"])
    hits = [w for w in meta.delivery_warnings if "四、存货" in w]
    assert hits, meta.delivery_warnings
    assert not [w for w in hits if "四、存货 存货" in w], hits


def test_outline_regions_carry_structure_titles_including_trailing_captions():
    from app.services.note_deliverable_outline import load_note_template_outline

    ol = load_note_template_outline("soe", "standalone")
    if ol is None:
        pytest.skip("真实附注模板缺失")
    by_code = {r.code: r for r in ol.regions if r.code}
    cash, intangible = by_code["八、1"], by_code["八、27"]
    assert cash.filter_titles == ("财务报表主要项目注释", "货币资金")
    assert cash.structure == ("受限制的货币资金明细",)  # 块结束标记之后的块尾表题归本块
    assert intangible.structure == ("无形资产情况", "确认为无形资产的数据资源")
    policy = next(r for r in ol.regions if r.code == "四、无形资产")
    assert not [t for t in policy.structure if t.endswith(("：", ":"))]  # 「…摊销方法如下：」是引导句，不当结构标题


@pytest.mark.asyncio
async def test_old_import_text_only_repeating_template_captions_is_not_delivered():
    """华氏旧导入：正文只有模板表题 +「本期无此项业务。」→ 交付件不出具、报表附注列不编号，并提醒报表有余额。"""
    from app.services.note_deliverable_outline import build_fs_note_ref_plan, load_note_template_outline
    from app.services.note_word_exporter import NoteWordExporter

    if load_note_template_outline("soe", "standalone") is None:
        pytest.skip("真实附注模板缺失")
    junk = _note("八、27", "无形资产", text="无形资产情况\n本期无此项业务。\n确认为无形资产的数据资源\n本期无此项业务。")
    notes = [*_real_notes(), junk]
    rows = [*_FS_ROWS, ("无形资产", 1000.0, 1200.0)]
    buf, meta = await NoteWordExporter(_FakeDB(notes, rows)).export_with_meta(
        SimpleNamespace(), 2025, template_type="soe", report_scope="standalone", mode="template")
    plan = await build_fs_note_ref_plan(_FakeDB(notes, rows), SimpleNamespace(), 2025,
                                        template_type="soe", report_scope="standalone")
    assert "八、27" not in meta.kept_codes and plan.ref_for("无形资产") is None
    assert not [t for t in _texts(Document(buf)) if "无形资产情况" in t]
    assert any("报表项目有余额" in w and "无形资产" in w for w in meta.delivery_warnings)


@pytest.mark.asyncio
async def test_sub_item_section_keeps_its_parent_title_and_numbering_matches():
    """国企模板「一年内到期的非流动负债」分项「（1）一年内到期的长期借款」是单独的块、标题取消了自动编号：
    分项有数而上级没有时上级只留标题；分项不占节号（报表附注列与交付件同口径）。"""
    from app.services.note_deliverable_outline import build_fs_note_ref_plan, load_note_template_outline
    from app.services.note_word_exporter import NoteWordExporter

    ol = load_note_template_outline("soe", "standalone")
    if ol is None:
        pytest.skip("真实附注模板缺失")
    by_code = {r.code: r for r in ol.regions if r.code}
    assert by_code["八、44"].numbered and not by_code["八、45"].numbered
    notes = [*_real_notes(),
             _note("八、45", "（1）一年内到期的长期借款", tables=[_table("一年内到期的长期借款", "信用借款", 30.0)])]
    rows = [*_FS_ROWS, ("一年内到期的非流动负债", 30.0, 0.0)]
    buf, meta = await NoteWordExporter(_FakeDB(notes, rows)).export_with_meta(
        SimpleNamespace(), 2025, template_type="soe", report_scope="standalone", mode="template")
    plan = await build_fs_note_ref_plan(_FakeDB(notes, rows), SimpleNamespace(), 2025,
                                        template_type="soe", report_scope="standalone")
    texts = [t for t in _texts(Document(buf)) if t]
    i = texts.index("一年内到期的非流动负债")
    assert texts[i + 1] == "（1）一年内到期的长期借款" and "信用借款" in texts[i + 2:i + 12]
    assert plan.ref_for("一年内到期的非流动负债") == "七、3"  # 货币资金 / 衍生金融资产之后
    assert plan.numbers.get(by_code["八、45"].order) is None
    assert {"八、44", "八、45"} <= set(meta.kept_codes)

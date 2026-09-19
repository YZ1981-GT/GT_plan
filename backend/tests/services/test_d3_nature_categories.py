"""D3「款项性质 / 关联方类型」枚举与源模板三向锁死守卫。

Spec: .kiro/specs/d-cycle-four-table-extraction-and-disclosure-completion/ Task 26
Validates: Requirements 10.1, 10.2, 10.3

判据真源 = 源 xlsx（openpyxl 直读，data_only=False）：
    backend/wp_templates/D/D3 预收账款.xlsx
        审定表D3-1!A8:A11                  四类款项性质（SUMIF 匹配键）
        审定表D3-1!B8                      SUMIF('预收账款明细表D3-2'!$C$12:$C$22, A8, ...)
        预收账款明细表D3-2!C12:C22 DV       同四类性质
        预收账款明细表D3-2!D12:D23 DV       三类关联方类型
        预收账款明细表D3-2!A10             '对方单位名称'（⇒ 行维度是对方单位，不是科目）

与之交叉锁死的前端单一真源：
    audit-platform/frontend/src/components/workpaper/composables/d3NatureCategories.ts

不连库，可进 CI。
"""
from __future__ import annotations

import re
from pathlib import Path

import openpyxl
import pytest

# ─── 路径定位（双哨兵向上查找）──────────────────────────────────────────────


def _find_repo_root() -> Path:
    d = Path(__file__).resolve()
    for _ in range(12):
        d = d.parent
        if (d / "backend" / "app" / "main.py").exists() and (
            d / "audit-platform" / "frontend" / "package.json"
        ).exists():
            return d
    raise AssertionError(f"repo root not found from {Path(__file__).resolve()}")


REPO_ROOT = _find_repo_root()
SRC_XLSX = REPO_ROOT / "backend" / "wp_templates" / "D" / "D3 预收账款.xlsx"
FE_TS = (
    REPO_ROOT
    / "audit-platform"
    / "frontend"
    / "src"
    / "components"
    / "workpaper"
    / "composables"
    / "d3NatureCategories.ts"
)
D3_RENDER = (
    REPO_ROOT / "backend" / "app" / "routers" / "wp_render_strategies" / "_d3_prepaid_accounts.py"
)

SHEET_ADJ = "审定表D3-1"
SHEET_DETAIL = "预收账款明细表D3-2"

EXPECTED_NATURES = [
    "预收销售固定资产款",
    "预收销售土地使用权款",
    "合同不成立时已收取的对价",
    "其他",
]
EXPECTED_RELATIONS = ["合并范围内关联方", "合并范围外关联方", "非关联方"]

# 源模板 C23 单独挂的另一套枚举（缺陷①，禁被实现成候选）
DEFECT_C23_ENUM = [
    "货款",
    "工程款",
    "设备款",
    "服务费",
    "建造合同形成的已结算尚未完工款",
    "其他",
]


# ─── fixtures ───────────────────────────────────────────────────────────────


@pytest.fixture(scope="module")
def wb():
    assert SRC_XLSX.exists(), f"源模板缺失: {SRC_XLSX}"
    return openpyxl.load_workbook(SRC_XLSX, data_only=False)


@pytest.fixture(scope="module")
def fe_src() -> str:
    assert FE_TS.exists(), f"前端真源缺失: {FE_TS}"
    return FE_TS.read_text(encoding="utf-8")


def _dv_values(ws, coord_probe: str) -> list[str] | None:
    """取覆盖 coord_probe 的 list 型数据验证取值域（按 sqref 逐格测试，禁按打印顺序取第一个）。"""
    for dv in ws.data_validations.dataValidation:
        if dv.type != "list":
            continue
        if coord_probe not in dv.sqref:
            continue
        f1 = (dv.formula1 or "").strip()
        if f1.startswith('"') and f1.endswith('"'):
            return [s for s in f1[1:-1].split(",") if s]
    return None


def _array_body(src: str, const_name: str) -> str:
    """截出 `export const X<类型注解> = ...[ ... ]` 的数组体。

    🔴 必须从赋值号 `=` 之后开始找 `[`：声明形如
    `export const X: readonly T[] = Object.freeze([...])`，
    直接找 anchor 之后第一个 `[` 会命中**类型注解 `T[]`** 的括号，
    body 退化成 `[]`、抽取恒空 → 上层断言变成空集比较（假红）。
    """
    at = src.index(f"export const {const_name}")
    eq = src.index("=", at)
    start = src.index("[", eq)
    depth = 0
    end = -1
    for i in range(start, len(src)):
        if src[i] == "[":
            depth += 1
        elif src[i] == "]":
            depth -= 1
            if depth == 0:
                end = i
                break
    assert end > start, f"{const_name} 数组体括号未配对"
    return src[start : end + 1]


def _fe_labels(src: str) -> list[str]:
    return re.findall(r"label:\s*'([^']+)'", _array_body(src, "D3_NATURE_CATEGORIES"))


def _fe_relations(src: str) -> list[str]:
    return re.findall(r"'([^']+)'", _array_body(src, "D3_RELATION_TYPES"))


# ─── Property 32: 源模板性质枚举 ────────────────────────────────────────────


class TestSourceTemplateNatures:
    def test_adjudication_nature_rows_are_four(self, wb):
        ws = wb[SHEET_ADJ]
        assert ws["A5"].value == "一、按照性质分类"
        got = [ws[f"A{r}"].value for r in range(8, 12)]
        assert got == EXPECTED_NATURES, f"审定表D3-1!A8:A11 变了: {got}"

    def test_row12_is_expandable_and_row13_is_total(self, wb):
        """R12 是空可扩行且计入合计（合计 SUM 覆盖 B8:B12）。"""
        ws = wb[SHEET_ADJ]
        assert ws["A12"].value is None, "R12 应为空可扩行"
        assert ws["A13"].value == "合计"
        assert ws["B13"].value == "=SUM(B8:B12)", ws["B13"].value

    def test_sumif_matches_detail_c_column(self, wb):
        """性质行金额 = 按 D3-2 C 列（款项性质）SUMIF ⇒ 该列是联动键。"""
        ws = wb[SHEET_ADJ]
        f = ws["B8"].value or ""
        assert f.startswith("=SUMIF("), f
        assert f"'{SHEET_DETAIL}'!$C$12:$C$22" in f, f
        assert f"'{SHEET_ADJ}'!A8" in f, f

    def test_detail_c_column_dv_equals_natures(self, wb):
        ws = wb[SHEET_DETAIL]
        vals = _dv_values(ws, "C12")
        assert vals is not None, "D3-2!C12 无 list 型数据验证"
        assert vals == EXPECTED_NATURES, vals

    def test_detail_row_dimension_is_counterparty_not_account(self, wb):
        """D3-2 行维度是对方单位（客户），不是 2203 子科目 —— requirements 10.1 的判据。"""
        ws = wb[SHEET_DETAIL]
        assert ws["A10"].value == "对方单位名称", ws["A10"].value
        assert ws["B10"].value == "公司代码"
        assert ws["C10"].value == "款项性质"
        assert ws["D10"].value == "关联方类型"

    def test_detail_d_column_dv_equals_relations(self, wb):
        ws = wb[SHEET_DETAIL]
        vals = _dv_values(ws, "D12")
        assert vals is not None, "D3-2!D12 无 list 型数据验证"
        assert vals == EXPECTED_RELATIONS, vals


# ─── Property 33: 前后端交叉锁死 ────────────────────────────────────────────


class TestFrontendCrossLock:
    def test_frontend_natures_equal_source(self, wb, fe_src):
        ws = wb[SHEET_ADJ]
        src_labels = [ws[f"A{r}"].value for r in range(8, 12)]
        assert _fe_labels(fe_src) == src_labels

    def test_frontend_relations_equal_source_dv(self, wb, fe_src):
        vals = _dv_values(wb[SHEET_DETAIL], "D12")
        assert _fe_relations(fe_src) == vals

    def test_frontend_source_refs_point_to_real_cells(self, wb, fe_src):
        """每个 sourceRef 指向的单元格必须真含该 label。"""
        ws = wb[SHEET_ADJ]
        pairs = re.findall(r"label:\s*'([^']+)',\s*sourceRef:\s*'([^']+)'", fe_src)
        assert len(pairs) == 4, pairs
        for label, ref in pairs:
            sheet, coord = ref.split("!")
            assert sheet == SHEET_ADJ, ref
            assert ws[coord].value == label, f"{ref} 实为 {ws[coord].value!r}，声明为 {label!r}"

    def test_frontend_relation_source_ref_literal(self, fe_src):
        assert f"'{SHEET_DETAIL}!D12:D23'" in fe_src

    def test_extraction_helpers_not_vacuous(self, fe_src):
        """反向自检：抽取器非空转（否则上面几条断言全是空集比较）。"""
        assert len(_fe_labels(fe_src)) == 4
        assert len(_fe_relations(fe_src)) == 3


# ─── Property 34: 源模板缺陷登记 ────────────────────────────────────────────


class TestSourceTemplateDefects:
    def test_defect1_c23_has_foreign_enum(self, wb):
        """缺陷①：C23 单独挂另一套 6 项枚举（与 D3-1 SUMIF 匹配的四类完全不同）。"""
        ws = wb[SHEET_DETAIL]
        vals = _dv_values(ws, "C23")
        assert vals is not None, "C23 无 list 型 DV（源模板已变，需复核缺陷登记）"
        assert vals == DEFECT_C23_ENUM, vals
        assert vals != EXPECTED_NATURES

    def test_defect1_foreign_enum_not_implemented(self, fe_src):
        """那套错误枚举不得被实现成平台候选。"""
        labels = _fe_labels(fe_src)
        for wrong in DEFECT_C23_ENUM:
            if wrong == "其他":  # 两套枚举共有的合法值
                continue
            assert wrong not in labels, f"源模板 C23 的错误枚举 {wrong} 被实现成候选"

    def test_defect2_sumif_range_shorter_than_total_range(self, wb):
        """缺陷②：SUMIF 只到 R22，而合计 E24=SUM(E12:E23) 含 R23 ⇒ R23 录入进合计不进性质行。"""
        adj = wb[SHEET_ADJ]
        det = wb[SHEET_DETAIL]
        assert "$C$12:$C$22" in (adj["B8"].value or "")
        assert det["A24"].value == "合计"
        assert det["E24"].value == "=SUM(E12:E23)", det["E24"].value

    def test_defects_registered_in_frontend(self, fe_src):
        """两处缺陷必须在前端真源里显式登记（防后来者照抄源模板）。

        🔴 判据必须逐个比对 `ref:` **字段值**，不能用「整段包含 'C23'」——
        第一条缺陷的 note 正文里也写了 C23，改坏 ref 值仍会通过（变异检验实测 GREEN）。
        """
        block = _array_body(fe_src, "D3_SOURCE_TEMPLATE_DEFECTS")
        refs = re.findall(r"ref:\s*(?:'([^']+)'|\"([^\"]+)\")", block)
        ref_values = [a or b for a, b in refs]
        assert len(ref_values) >= 2, f"缺陷登记条目少于 2: {ref_values}"
        assert f"{SHEET_DETAIL}!C23" in ref_values, ref_values
        assert any("SUMIF" in r for r in ref_values), ref_values

        # 每条 note 说明须实质（≥60 字），防退化成占位文字
        entries = re.findall(r"ref:[\s\S]*?note:([\s\S]*?)\}\)", block)
        assert len(entries) == len(ref_values), (len(entries), len(ref_values))
        for ref, note in zip(ref_values, entries):
            text = "".join(re.findall(r"'([^']*)'", note))
            assert len(text) >= 60, f"{ref} 的 note 过短({len(text)}字): {text[:40]}"


# ─── Property 35: 反向锁死禁按 2203 子科目建 D3-2 行 ────────────────────────


class TestNoAccountKeyedDetailRows:
    """requirements 10.1 曾写「D3-2 按 2203 叶子子科目建行」，按源模板不成立。

    D3-2 行维度是对方单位（`A10='对方单位名称'`），数据来源是 `tb_aux_balance` 2203
    **客户维度**归集（render 已记宁缺勿造决策）。故 render 不得产出按科目名建的明细行。
    """

    @staticmethod
    def _strip_py_comments(src: str) -> str:
        out: list[str] = []
        for line in src.splitlines():
            s = line.split("#", 1)[0]
            out.append(s)
        joined = "\n".join(out)
        # 粗略剥三引号 docstring
        return re.sub(r'"""[\s\S]*?"""', "", joined)

    def test_render_has_no_account_keyed_detail_prefill(self):
        assert D3_RENDER.exists(), D3_RENDER
        src = self._strip_py_comments(D3_RENDER.read_text(encoding="utf-8"))
        for bad in ("预收货款", "预收项目款"):
            assert bad not in src, f"D3 render 出现 2203 子科目名 {bad}（D3-2 行维度是对方单位）"
        assert "detail_prefill" not in src, "D3 render 不应产出明细预填（宁缺勿造决策在案）"

    def test_render_records_ningquewuzao_decision(self):
        """决策必须留在源码（含 docstring/注释），否则下个会话会重新实现一遍。"""
        raw = D3_RENDER.read_text(encoding="utf-8")
        assert "宁缺勿造" in raw
        assert "tb_aux_balance" in raw
        assert "客户维度" in raw

    def test_strip_helper_self_check(self):
        """反向自检：剥注释确实生效。

        🔴 锚点不能用「宁缺勿造」—— 它**也出现在 `logger.debug(...)` 字符串里**（非注释），
        剥注释后仍在 ⇒ 那样的自检恒失败且掩盖真实能力。改用只在 `#` 注释块里出现的
        `【宁缺勿造决策】`，并另断言字符串里的那份确实被保留（证明只剥了注释）。
        """
        raw = D3_RENDER.read_text(encoding="utf-8")
        stripped = self._strip_py_comments(raw)
        assert "【宁缺勿造决策】" in raw
        assert "【宁缺勿造决策】" not in stripped, "剥注释未生效"
        assert "宁缺勿造" in stripped, "logger 字符串里的那份不应被剥掉（证明只剥注释不剥字符串）"
        assert len(stripped) > len(raw) * 0.2

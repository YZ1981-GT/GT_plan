"""Word 导出复用投影器：导出表结构（表名/列头/行）== 投影结果（P12）

spec: disclosure-table-sync-convergence Task 5.2
仅测试 NoteWordExporter._note_tables / _effective_table_data 的投影集成，
不生成完整 docx（表结构一致即证明导出与模块渲染同源）。
"""

from __future__ import annotations

from types import SimpleNamespace

from app.services.note_word_exporter import NoteWordExporter
from app.services.note_sub_table_projector import project_sub_tables


def _exporter() -> NoteWordExporter:
    return NoteWordExporter(db=None)  # _note_tables/_effective_table_data 不触库


def _workpaper_note():
    td = {
        "_source": "workpaper",
        "sub_table_data": {
            "存货分类": [
                {"label": "原材料", "end_gross": 100, "end_impairment": 5},
                {"label": "合计", "end_gross": 100, "end_impairment": 5, "is_total": True},
            ]
        },
        "_sub_table_columns": {
            "存货分类": [
                {"key": "label", "label": "存货类别", "is_label": True},
                {"key": "end_gross", "label": "期末账面余额"},
                {"key": "end_impairment", "label": "期末跌价准备"},
            ]
        },
    }
    return SimpleNamespace(table_data=td, text_content=None)


def test_p12_export_tables_equal_projection():
    """workpaper 来源：_note_tables 返回的表 == project_sub_tables 结果（同源）。"""
    note = _workpaper_note()
    exporter = _exporter()

    projected = project_sub_tables(note.table_data)
    export_tables = exporter._note_tables(note)

    assert projected is not None and len(projected) == 1
    assert len(export_tables) == 1
    et, pt = export_tables[0], projected[0]
    assert et["name"] == pt["name"] == "存货分类"
    assert et["headers"] == pt["headers"] == ["存货类别", "期末账面余额", "期末跌价准备"]
    assert et["rows"] == pt["rows"]
    # 合计行保持
    assert et["rows"][-1]["is_total"] is True


def test_p12_has_content_true_for_workpaper_projected():
    """workpaper 投影出非零值 → _has_content 为 True（不被误判空章节跳过）。"""
    exporter = _exporter()
    note = SimpleNamespace(
        table_data=_workpaper_note().table_data,
        text_content=None,
        is_deleted=False,
        status=None,
    )
    assert exporter._has_content(note) is True


def test_engine_source_export_unchanged():
    """非 workpaper 来源：_note_tables 不投影，沿用既有 _tables（导出行为不变，Req6.2）。"""
    exporter = _exporter()
    td = {
        "_source": "engine_fill",
        "_tables": [
            {"name": "既有表", "headers": ["项目", "金额"],
             "rows": [{"label": "行1", "values": [10]}], "export_enabled": True}
        ],
    }
    note = SimpleNamespace(table_data=td, text_content=None)
    tables = exporter._note_tables(note)
    assert len(tables) == 1
    assert tables[0]["name"] == "既有表"
    assert tables[0]["headers"] == ["项目", "金额"]


def test_export_enabled_false_filtered_after_projection():
    """投影表默认 export_enabled 缺省视为 True（不被过滤）。"""
    exporter = _exporter()
    note = _workpaper_note()
    tables = exporter._note_tables(note)
    assert len(tables) == 1  # 投影表无 export_enabled=false → 保留


# ---------------------------------------------------------------------------
# _build_two_level_header_rows：row1 不得为 rowspan 列补占位
#
# 回归来源：Word 导出附注存货章节实测发现第二行子表头整体右移一格、末列丢失
# （「存货分类」少末尾「账面价值」；「按组合计提」少第二个「比例(%)」）。
# 根因：无分组列已用 rowspan=2 纵向合并，row1 再补空占位会挤占一个真实列位，
# 而 fill_multi_header 本就会跳过被 rowspan 占用的列。
#
# spec: f2-inventory-disclosure-template-alignment R11.2
# ---------------------------------------------------------------------------

INVENTORY_HEADERS = [
    "项目",
    "账面余额", "跌价准备/合同履约成本减值准备", "账面价值",
    "账面余额", "跌价准备/合同履约成本减值准备", "账面价值",
]
INVENTORY_GROUPS = [
    {"group": "期末余额", "start": 1, "span": 3},
    {"group": "上年年末余额", "start": 4, "span": 3},
]


def test_build_two_level_header_rows_row1_has_only_subcolumns() -> None:
    """row1 长度 = 分组覆盖的列数（不含 rowspan 列占位）。"""
    from app.services.note_word_exporter import _build_two_level_header_rows

    row0, row1 = _build_two_level_header_rows(INVENTORY_HEADERS, INVENTORY_GROUPS)

    assert [c["text"] for c in row0] == ["项目", "期末余额", "上年年末余额"]
    assert row0[0]["rowspan"] == 2
    assert row0[1]["colspan"] == 3 and row0[2]["colspan"] == 3

    # 6 个子列名，且不含空串占位
    assert [c["text"] for c in row1] == [
        "账面余额", "跌价准备/合同履约成本减值准备", "账面价值",
        "账面余额", "跌价准备/合同履约成本减值准备", "账面价值",
    ]
    assert all(c["text"] for c in row1), "row1 不得含空占位（会挤掉末列）"


def test_two_level_header_renders_all_subcolumns_in_docx() -> None:
    """端到端：渲染进 docx 后每个子列名都在正确列位，末列不丢。"""
    from docx import Document

    from app.services.note_word_exporter import (
        _build_two_level_header_rows,
        fill_multi_header,
    )

    header_rows = _build_two_level_header_rows(INVENTORY_HEADERS, INVENTORY_GROUPS)
    doc = Document()
    table = doc.add_table(rows=2, cols=len(INVENTORY_HEADERS))
    fill_multi_header(table, header_rows, total_cols=len(INVENTORY_HEADERS))

    r1 = [c.text.strip() for c in table.rows[1].cells]
    # 索引 0 是标签列的纵向合并延续（python-docx 回显主格文本）
    assert r1[1:] == [
        "账面余额", "跌价准备/合同履约成本减值准备", "账面价值",
        "账面余额", "跌价准备/合同履约成本减值准备", "账面价值",
    ], f"子列名错位或丢失：{r1}"


def test_two_level_header_multi_group_with_trailing_plain_column() -> None:
    """「按组合计提」形态：分组之后还有无分组列（账面价值）也不错位。"""
    from docx import Document

    from app.services.note_word_exporter import (
        _build_two_level_header_rows,
        fill_multi_header,
    )

    headers = ["组合", "金额", "比例(%)", "金额", "计提标准", "比例(%)", "账面价值"]
    groups = [
        {"group": "账面余额", "start": 1, "span": 2},
        {"group": "存货跌价准备", "start": 3, "span": 3},
    ]

    row0, row1 = _build_two_level_header_rows(headers, groups)
    assert [c["text"] for c in row0] == ["组合", "账面余额", "存货跌价准备", "账面价值"]
    assert [c["text"] for c in row1] == ["金额", "比例(%)", "金额", "计提标准", "比例(%)"]

    doc = Document()
    table = doc.add_table(rows=2, cols=len(headers))
    fill_multi_header(table, [row0, row1], total_cols=len(headers))
    r1 = [c.text.strip() for c in table.rows[1].cells]
    # 第二个「比例(%)」曾因错位丢失
    assert r1[1:6] == ["金额", "比例(%)", "金额", "计提标准", "比例(%)"], f"错位：{r1}"


def test_all_plain_columns_yields_empty_row1() -> None:
    """全无分组时 row1 为空（调用方据此判定不走两行表头）。"""
    from app.services.note_word_exporter import _build_two_level_header_rows

    row0, row1 = _build_two_level_header_rows(["项目", "期初", "期末"], [])
    assert [c["text"] for c in row0] == ["项目", "期初", "期末"]
    assert row1 == []

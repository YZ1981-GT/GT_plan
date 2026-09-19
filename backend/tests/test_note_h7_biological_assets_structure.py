"""附注 H7 生产性生物资产章节结构守卫（上市 §五、24 / 国企 §八、24）。

锁定 `fix_note_h7_biological_assets_structure.py` 的重建结果。H7 是 H 循环唯一需要
**整体重建**的循环（重建前两个披露 Tab 只有单行只读 movementRows，无同步链路）。

固化三条最易踩错的判断：

1. **上市两表是列转置 + 两级表头**（openpyxl 合并区 `A9:A10`/`B9:C9`/…/`J9:J10` 实证）
   → `group` = 产业名；项目列与合计列不带 `group`（rowspan=2）；**不得标 `flat`**。
2. **`……` 两种语义相反**：作**列头**（R10/R52）永远收不到数据 → 丢弃；
   作**行**（上市表1 R20/R30/R40、表2 R58/R63）是真实可扩明细行且参与小计 → **保留**。
   国企的 `……`（R11/R14/R17/R20）是行形态的「可继续加行」标记，动态类别行模型下不需要
   → seed 删除（13 行 → 9 行）。
3. **上市列 key 不能用 label**：四个产业的默认叶子名都是源模板字面 `类别`，会撞键
   → 用稳定 `{industryKey}_{seq}`。

spec: h7-biological-assets-disclosure-rebuild (Task 8)
"""
from __future__ import annotations

import importlib.util
import json
import re
from pathlib import Path

import openpyxl

_ROOT = Path(__file__).resolve().parent.parent
_FIX = _ROOT / "scripts" / "fix" / "fix_note_h7_biological_assets_structure.py"
_XLSX = _ROOT / "wp_templates" / "H" / "H7 生产性生物资产.xlsx"
_FE = (
    _ROOT.parent / "audit-platform" / "frontend" / "src" / "components"
    / "workpaper" / "composables"
)
_FE_MAP = _FE / "h7NoteSectionMap.ts"
_FE_PAYLOAD = _FE / "h7DisclosureSyncPayload.ts"
_FE_LISTED_MODEL = _FE / "h7ListedDisclosureModel.ts"
_FE_SOE_MODEL = _FE / "h7SoeDisclosureModel.ts"
_MATRIX = _ROOT / "data" / "note_template_variant_matrix.json"

_SHEET_LISTED = "附注披露信息（上市公司）"
_SHEET_SOE = "附注披露信息（国有企业）"
_LEAKED = "项  目"
_ELLIPSIS = "……"
_INDUSTRY_LABELS = ["种植业", "畜牧养殖业", "林业", "水产业"]


def _load_fix():
    spec = importlib.util.spec_from_file_location("_fix_h7_struct", _FIX)
    mod = importlib.util.module_from_spec(spec)
    spec.loader.exec_module(mod)
    return mod


FIX = _load_fix()


def _norm(v: object) -> str:
    return re.sub(r"\s+", "", str(v or ""))


def _strip_ts_comments(src: str) -> str:
    """去掉 TS 注释再做源码断言（注释里会写反例说明，不剥离会被数成真实代码）。"""
    src = re.sub(r"/\*[\s\S]*?\*/", "", src)
    return re.sub(r"(^|[^:])//.*$", r"\1", src, flags=re.M)


def _section(key: str) -> dict:
    path, section_number, _plan, _ts = FIX._TARGETS[key]
    doc = json.loads(Path(path).read_text(encoding="utf-8"))
    sec = next(
        (s for s in doc["sections"] if str(s.get("section_number")) == section_number), None
    )
    assert sec is not None, f"{key} 章节 {section_number!r} 不存在"
    return sec


def _tables(key: str) -> dict[str, dict]:
    return {str(t.get("name")): t for t in (_section(key).get("tables") or [])}


def _ws(sheet: str):
    wb = openpyxl.load_workbook(_XLSX, data_only=True)
    assert sheet in wb.sheetnames, wb.sheetnames
    return wb[sheet]


def _src_col1(sheet: str, first: int, last: int) -> list[str]:
    ws = _ws(sheet)
    return [_norm(ws.cell(r, 1).value) for r in range(first, last + 1)]


# ───────────────────── Property 10：幂等与零欠账 ─────────────────────


def test_check_passes_both_variants():
    for key in FIX._TARGETS:
        _changes, warnings, errs = FIX._runner(key, dry_run=True, check=True)
        assert not errs, f"{key} 结构欠账：{errs}"
        assert not warnings, f"{key} 告警：{warnings}"


def test_templates_parse():
    for path in {t[0] for t in FIX._TARGETS.values()}:
        json.loads(Path(path).read_text(encoding="utf-8"))


def test_variant_matrix_sections():
    matrix = json.loads(_MATRIX.read_text(encoding="utf-8"))
    acct = next(
        a for a in matrix["accounts"] if a["account_key"] == "sheng_chan_xing_sheng_wu_zi_chan"
    )
    assert acct["variants"]["listed_standalone"] == FIX.LISTED_SECTION
    assert acct["variants"]["soe_standalone"] == FIX.SOE_SECTION


# ───────── Property 2：上市两级表头与源模板合并区一致 ─────────


def test_listed_source_merged_ranges_prove_two_level_header():
    """源 xlsx 合并区实证两级表头（改坏了就说明源模板判读依据变了，需重新确权）。"""
    ws = _ws(_SHEET_LISTED)
    merged = {str(m) for m in ws.merged_cells.ranges}
    for rng in ("A9:A10", "B9:C9", "D9:E9", "F9:G9", "H9:I9", "J9:J10"):
        assert rng in merged, f"表1 缺合并区 {rng}：{sorted(merged)}"
    for rng in ("A51:A52", "B51:C51", "D51:E51", "F51:G51", "H51:I51", "J51:J52"):
        assert rng in merged, f"表2 缺合并区 {rng}"
    # R9 父表头 = 4 产业 + 合计；R10 子表头 = `类别 | ……` × 4
    assert [_norm(ws.cell(9, c).value) for c in (2, 4, 6, 8, 10)] == [
        *_INDUSTRY_LABELS, "合计",
    ]
    assert [_norm(ws.cell(10, c).value) for c in range(2, 10)] == [
        "类别", _ELLIPSIS, "类别", _ELLIPSIS, "类别", _ELLIPSIS, "类别", _ELLIPSIS,
    ]


def test_listed_columns_two_level():
    for name in (FIX.T_LISTED_COST, FIX.T_LISTED_FAIR):
        tbl = _tables("listed")[name]
        cols = tbl.get("columns") or []
        assert [c.get("key") for c in cols] == [
            "label", "crop_1", "livestock_1", "forestry_1", "aquatic_1", "total",
        ], (name, cols)
        assert [c.get("label") for c in cols] == tbl["headers"], name
        assert [c.get("label") for c in cols] == [
            "项目", "类别", "类别", "类别", "类别", "合计",
        ], name
        assert cols[0].get("is_label") is True, name
        assert not cols[0].get("group"), f"{name} 项目列 rowspan=2 不得带 group"
        assert not cols[-1].get("group"), f"{name} 合计列 rowspan=2 不得带 group"
        assert [c.get("group") for c in cols[1:-1]] == _INDUSTRY_LABELS, name
        assert not any(c.get("flat") for c in cols), f"{name} 两级表头不得标 flat"
        assert all(c.get("format") == "amount" for c in cols[1:]), name
        assert tbl.get("_column_groups") == [
            {"group": lb, "start": i + 1, "span": 1} for i, lb in enumerate(_INDUSTRY_LABELS)
        ], (name, tbl.get("_column_groups"))


def test_ellipsis_never_a_column():
    """`……` 作列头永远收不到数据 → 两版列定义与 headers 都不得含它。"""
    for key in FIX._TARGETS:
        for name, tbl in _tables(key).items():
            assert _ELLIPSIS not in tbl["headers"], f"{key}/{name} headers 残留 `……` 列头"
            labels = [c.get("label") for c in tbl.get("columns") or []]
            assert _ELLIPSIS not in labels, f"{key}/{name} columns 残留 `……` 列头"


def test_listed_column_keys_are_not_labels():
    """四个产业默认叶子名都是 `类别` → 列 key 必须是稳定标识，不能用 label。"""
    cols = _tables("listed")[FIX.T_LISTED_COST]["columns"]
    keys = [c["key"] for c in cols]
    assert len(set(keys)) == len(keys), f"列 key 撞键：{keys}"
    assert "类别" not in keys


# ───────── Property 1：上市行集与源 xlsx 一致 + `……` 行保留 ─────────


def test_listed_cost_rows_match_source_xlsx():
    rows = _tables("listed")[FIX.T_LISTED_COST].get("rows") or []
    tpl = [_norm(r.get("label")) for r in rows]
    src = _src_col1(_SHEET_LISTED, 11, 44)
    assert len(src) == 34, len(src)
    assert tpl == src, f"上市表1 行集漂移\n模板={tpl}\n源={src}"


def test_listed_cost_keeps_three_ellipsis_rows():
    """🔴 正向断言：`……` 作行是真实可扩明细行（参与本期减少小计），不得被当占位删除。"""
    rows = _tables("listed")[FIX.T_LISTED_COST].get("rows") or []
    at = [i for i, r in enumerate(rows) if _norm(r.get("label")) == _ELLIPSIS]
    assert len(at) == 3, f"上市表1 `……` 可扩行应为 3 行（三层各一处本期减少），实测 {at}"
    assert at == [9, 19, 29], at


def test_listed_fair_rows_match_source_xlsx():
    rows = _tables("listed")[FIX.T_LISTED_FAIR].get("rows") or []
    tpl = [_norm(r.get("label")) for r in rows]
    # 源 R53–R64，其中 R61 源模板留白
    src = [x for x in _src_col1(_SHEET_LISTED, 53, 64) if x]
    assert len(src) == 11, len(src)
    assert tpl == src, f"上市表2 行集漂移\n模板={tpl}\n源={src}"
    at = [i for i, r in enumerate(rows) if _norm(r.get("label")) == _ELLIPSIS]
    assert at == [5, 9], f"上市表2 应保留 2 个 `……` 可扩行，实测 {at}"


# ───────── Property 5/7：国企行集 9 行、表名正名 ─────────


def test_soe_columns_flat():
    ws = _ws(_SHEET_SOE)
    for row_no, name in ((8, FIX.T_SOE_COST), (25, FIX.T_SOE_FAIR)):
        tbl = _tables("soe")[name]
        cols = tbl.get("columns") or []
        assert [c.get("key") for c in cols] == [
            "label", "begin", "increase", "decrease", "end",
        ], (name, cols)
        assert [c.get("label") for c in cols] == tbl["headers"], name
        assert cols[0].get("is_label") and cols[0].get("flat"), f"{name} 首列须 is_label + flat"
        assert not any(c.get("group") for c in cols), f"{name} 单级表头不得声明 group"
        assert not tbl.get("_column_groups"), f"{name} 单级表头不得残留 _column_groups"
        # 与源 xlsx 表头逐字比对
        assert [_norm(ws.cell(row_no, c).value) for c in range(1, 6)] == [
            "项目", "期初账面价值", "本期增加额", "本期减少额", "期末账面价值",
        ], name


def test_soe_rows_nine_and_no_ellipsis():
    """源 13 行含 4 个 `……`「可继续加行」标记 → 动态类别行模型下 seed 删除 → 9 行。"""
    for name in (FIX.T_SOE_COST, FIX.T_SOE_FAIR):
        rows = _tables("soe")[name].get("rows") or []
        labels = [_norm(r.get("label")) for r in rows]
        assert labels == [
            "一、种植业", "其中：1．",
            "二、畜牧养殖业", "其中：1．",
            "三、林业", "其中：1．",
            "四、水产业", "其中：1．",
            "合计",
        ], (name, labels)
        assert _ELLIPSIS not in labels, f"{name} 残留 `……` 占位行"
        assert rows[-1].get("is_total") is True, name


def test_soe_source_rows_still_have_ellipsis_markers():
    """反向确权：源模板确实有 4 个 `……` 行（本守卫的删除依据），源变了要重新判断。"""
    src = _src_col1(_SHEET_SOE, 9, 21)
    assert src.count(_ELLIPSIS) == 4, src
    assert [x for x in src if x.startswith("一、") or x.startswith("四、")] == [
        "一、种植业", "四、水产业",
    ]


def test_listed_table_renamed_and_leak_gone():
    names = list(_tables("listed"))
    assert names == [FIX.T_LISTED_COST, FIX.T_LISTED_FAIR], names
    assert _LEAKED not in names, "表头首格泄漏名复活"
    for key in FIX._TARGETS:
        assert _LEAKED not in _tables(key), f"{key} 残留 `{_LEAKED}`"


def test_no_header_label_fake_rows():
    for key in FIX._TARGETS:
        for name, tbl in _tables(key).items():
            assert not any(
                r.get("row_type") == "header_label" for r in tbl.get("rows") or []
            ), f"{key}/{name} 残留 header_label 假行（压扁的第二行表头）"


# ───────── Property 10：guidance / text_sections ─────────


def test_guidance_present():
    for key in FIX._TARGETS:
        for name, tbl in _tables(key).items():
            assert str(tbl.get("guidance") or "").strip(), f"{key}/{name} 缺 guidance"


def test_listed_text_sections():
    paras = [str(p) for p in (_section("listed").get("text_sections") or [])]
    assert f"#### {FIX.T_LISTED_COST}" in paras
    assert f"#### {FIX.T_LISTED_FAIR}" in paras
    assert any("公益性生物资产" in p for p in paras), "缺源模板 R7 公益性生物资产说明"
    assert any("15号文第十九条" in p for p in paras), "缺源模板 R46 减值披露要求"
    # 裸表名会被当披露正文渲染
    assert FIX.T_LISTED_COST not in paras and FIX.T_LISTED_FAIR not in paras


def test_soe_text_sections():
    paras = [str(p) for p in (_section("soe").get("text_sections") or [])]
    assert f"#### {FIX.T_SOE_COST}" in paras
    assert f"#### {FIX.T_SOE_FAIR}" in paras
    assert FIX.T_SOE_COST not in paras and FIX.T_SOE_FAIR not in paras
    # 源模板 R39 / R40 原本缺失或丢前缀
    assert "注：应披露公允价值确认依据。" in paras, "缺源模板 R39"
    assert any(p.startswith("（3）说明生产性生物资产相关的风险") for p in paras), "R40 缺「（3）」前缀"
    ws = _ws(_SHEET_SOE)
    assert _norm(ws.cell(39, 1).value) == "注：应披露公允价值确认依据。"
    assert _norm(ws.cell(40, 1).value).startswith("（3）说明生产性生物资产相关的风险")


# ───────── 前后端一致性（seed ↔ 推送同键） ─────────


def test_frontend_map_alignment():
    raw = _FE_MAP.read_text(encoding="utf-8")
    src = _strip_ts_comments(raw)
    assert f"listed: '{FIX.LISTED_SECTION}'" in src, "上市章节号漂移或改成了标识符引用"
    assert f"soe: '{FIX.SOE_SECTION}'" in src, "国企章节号漂移"
    # 🔴 H7 国企 tab 名是「国有企业」，与 H9/H10 的「国企」不同（连注释里也不许出现，防抄错）
    assert f"soe: '{_SHEET_SOE}'" in src, "国企 sheet 名未对齐源 xlsx tab 名"
    assert "附注披露信息（国企）" not in raw, "误用了 H9/H10 的「国企」写法"
    assert f"listed: '{_SHEET_LISTED}'" in src
    assert f"cost: '{FIX.T_LISTED_COST}'" in src and f"fair: '{FIX.T_LISTED_FAIR}'" in src
    assert f"cost: '{FIX.T_SOE_COST}'" in src and f"fair: '{FIX.T_SOE_FAIR}'" in src
    assert f"H7_LEGACY_OBSOLETE_TABLES: readonly string[] = ['{_LEAKED}']" in src


def test_frontend_note_section_object_body_has_no_comment():
    """`gen_note_wp_sync_registry.py` 从对象体抽 `listed: '…'`，体内注释会被优先抓到。"""
    src = _FE_MAP.read_text(encoding="utf-8")
    m = re.search(r"H7_NOTE_SECTION\s*=\s*\{(?P<body>[^}]*)\}", src, re.S)
    assert m, "未定位到 H7_NOTE_SECTION 对象"
    body = m.group("body")
    assert "//" not in body and "/*" not in body, f"对象体内含注释：{body!r}"
    assert body.count("listed") == 1, body


def test_frontend_payload_keys_align():
    src = _strip_ts_comments(_FE_PAYLOAD.read_text(encoding="utf-8"))
    for key in ("label", "begin", "increase", "decrease", "end"):
        assert f"key: '{key}'" in src, f"国企缺列 key {key}"
    assert "H7_NOTE_TEXT_TITLES" in src and "buildH7NoteTexts" in src
    assert "_removed_table_keys" in src
    assert "flat: true" in src, "国企单级表头须显式 flat"
    # 两级表头由 group 承载，不得给上市列标 flat
    assert "group: h7IndustryLabel" in src, "上市列须按产业分组"


def test_frontend_industry_parity():
    listed = _strip_ts_comments(_FE_LISTED_MODEL.read_text(encoding="utf-8"))
    block = listed.split("H7_INDUSTRIES", 1)[1].split("] as const", 1)[0]
    assert re.findall(r"label: '([^']+)'", block) == _INDUSTRY_LABELS, block
    soe = _strip_ts_comments(_FE_SOE_MODEL.read_text(encoding="utf-8"))
    sblock = soe.split("H7_SOE_INDUSTRIES", 1)[1].split("] as const", 1)[0]
    assert re.findall(r"label: '([^']+)'", sblock) == [
        "一、种植业", "二、畜牧养殖业", "三、林业", "四、水产业",
    ], sblock


def test_frontend_row_labels_match_template():
    """前端行定义标签序列 ≡ 模板行集（推送与 seed 同构）。"""
    listed = _strip_ts_comments(_FE_LISTED_MODEL.read_text(encoding="utf-8"))
    for const, name in (
        ("H7_COST_MOVEMENT_ROWS", FIX.T_LISTED_COST),
        ("H7_FAIR_MOVEMENT_ROWS", FIX.T_LISTED_FAIR),
    ):
        block = listed.split(f"{const}: H7MovementRowDef[] = [", 1)[1].split("\n]", 1)[0]
        labels = [_norm(x) for x in re.findall(r"label: '([^']+)'", block)]
        tpl = [_norm(r.get("label")) for r in _tables("listed")[name]["rows"]]
        assert labels == tpl, f"{const} 与模板 {name} 行集不一致\n前端={labels}\n模板={tpl}"


# ───────────────────── 反向自检（防守卫空转） ─────────────────────


def test_reverse_self_check_validate_section():
    from _note_structure_kit import validate_section

    broken = {
        "tables": [
            {"name": FIX.T_SOE_COST, "headers": ["项目"], "rows": [], "guidance": "x"}
        ]
    }
    errs = validate_section(broken, [FIX.T_SOE_COST])
    assert any("缺 columns" in e for e in errs), errs


def test_reverse_self_check_flat_group_conflict_detected():
    """两级表若误标 flat，validate_section 必须抓出表态冲突。"""
    from _note_structure_kit import validate_section

    bad = {
        "tables": [{
            "name": FIX.T_LISTED_COST,
            "headers": ["项目", "类别"],
            "columns": [
                {"key": "label", "label": "项目", "is_label": True, "flat": True},
                {"key": "crop_1", "label": "类别", "group": "种植业"},
            ],
            "rows": [],
            "guidance": "x",
        }]
    }
    errs = validate_section(bad, [FIX.T_LISTED_COST])
    assert any("flat 与 group 并存" in e for e in errs), errs


def test_reverse_self_check_comment_stripper():
    """`_strip_ts_comments` 既要剥掉注释、又不能吃掉代码（否则源码断言恒绿）。"""
    sample = (
        "/* 反例：listed: 'WRONG' */\n"
        "const a = 1 // listed: 'ALSO_WRONG'\n"
        "const url = 'https://x/y'\n"
        "const s = { listed: '五、24' }\n"
    )
    out = _strip_ts_comments(sample)
    assert "WRONG" not in out and "ALSO_WRONG" not in out
    assert "listed: '五、24'" in out, "把真实代码一起剥掉了"
    assert "https://x/y" in out, "不得把 URL 里的 // 当注释切掉"

    # 反向自检：映射源码的**注释**里确实写着 `listed: '…'` 形态的说明
    # → 证明上面对 `listed: '五、24'` 的断言依赖剥离，不是空转
    raw = _FE_MAP.read_text(encoding="utf-8")
    assert "listed: '…'" in raw, (
        "映射源码注释里已不再出现 `listed: '…'` 说明 → 请同步简化本反向自检"
    )
    assert "listed: '…'" not in _strip_ts_comments(raw)


def test_reverse_self_check_normalizer():
    assert _norm("合  计") == "合计"
    assert _norm("    1.期初余额") == "1.期初余额"
    assert _norm("   ……") == "……"
    assert _norm(None) == ""

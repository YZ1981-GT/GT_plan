"""调整分录导入导出 —— characterization（特征化）安全网。

spec: adjustment-import-export-contract / Task 1.1

目的：把「既有正确路径」的当前真实行为逐字节锁定，作为 Wave 0/1 改动的零回归对照。
本文件**只记录当前行为**，不表达"应该怎样"。凡当前行为可疑（后续波次会翻转）的断言，
均在注释中标明「Task X.Y 会翻转，翻转时同步更新本断言」。

锁定的基线：
  A. `_is_example_row` 对模板内置示例行的判定（Property 5 基线）
  B. 中央 `mode=append` 导入 `_import_adjustments` 的计数 / 分组合并 / 旧字段名兼容 /
     detail_account_code 判定，以及 `mode` 已传递到实现层的源码契约（Property 2 基线）
  C. 富模板 `GET /adjustments/export-template` 的 sheet 名集合与各 sheet 列头
     （后续 spec 明令不得改富模板列集合）+ `export-summary` 的 `_write_adj_sheet` 列头
  D. 7 张「无漂移」调整 sheet 的 item_id / 有效 storage_field / headers / field_keys
     （Property 13 基线）

不依赖真实 DB：用 fake AsyncSession（按调用顺序回放查询结果）+ patch AdjustmentService。
"""
from __future__ import annotations

import inspect
import re
from decimal import Decimal
from io import BytesIO
from types import SimpleNamespace
from uuid import uuid4

import openpyxl
import pytest

from app.models.audit_platform_models import AccountSource
from app.services.import_template_service import (
    TEMPLATE_COLUMNS,
    ImportType,
    _is_example_row,
)

# ═══════════════════════════════════════════════════════════════════════════
# 通用 fake DB（按 execute 调用顺序回放结果）
# ═══════════════════════════════════════════════════════════════════════════


class _FakeResult:
    def __init__(self, rows: list):
        self._rows = list(rows)

    def all(self):
        return list(self._rows)

    def scalar_one_or_none(self):
        return self._rows[0] if self._rows else None

    def scalars(self):
        return SimpleNamespace(all=lambda: list(self._rows))


class _FakeDB:
    """按顺序回放预置结果的最小 AsyncSession 替身（不触库）。"""

    def __init__(self, results: list[list]):
        self._queue = [_FakeResult(r) for r in results]
        self.execute_count = 0

    async def execute(self, *_a, **_kw):
        self.execute_count += 1
        if self._queue:
            return self._queue.pop(0)
        return _FakeResult([])

    async def flush(self):
        return None

    async def commit(self):
        return None

    async def rollback(self):
        return None

    def add(self, _obj):
        return None


# ═══════════════════════════════════════════════════════════════════════════
# A. `_is_example_row` —— 模板内置示例行判定基线
# ═══════════════════════════════════════════════════════════════════════════

_ADJ_COLUMNS = TEMPLATE_COLUMNS[ImportType.adjustments]

# 富模板（export-template）AJE模板 sheet 的两行内置示例（与 TEMPLATE_COLUMNS 示例值同源）：
#   ("AJE-001","AJE","补提应收账款减值","",<二级名称>,"","","",借,贷)
_BUILTIN_EXAMPLE_ROW_1 = [
    "AJE-001", "AJE", "补提应收账款减值", "", "应收账款", "", "", "", "0", "100000",
]
_BUILTIN_EXAMPLE_ROW_2 = [
    "AJE-001", "AJE", "补提应收账款减值", "", "信用减值损失", "", "", "", "100000", "0",
]


def test_a_matchable_pairs_and_threshold_baseline():
    """Task 1.3 新契约：判定为「与某个内置示例行的全部可比字段逐字相等」（无过半阈值）。

    可比字段 = 候选示例行中值非空的列；TEMPLATE_COLUMNS 单示例对 adjustments 是
    5 个（编号/类型/摘要/二级科目名称/贷方金额）。
    （原断言锁定的是「过半即示例」的旧宽松阈值，Task 1.3 已翻转 basis。）
    """
    matchable = [c for c in _ADJ_COLUMNS if str(c[3]).strip()]
    assert [c[0] for c in matchable] == ["编号", "类型", "摘要", "二级科目名称", "贷方金额"]

    src = inspect.getsource(_is_example_row)
    assert "match_count >= max(2, len(matchable_pairs) * 0.6)" not in src, (
        "宽松判定（过半即示例）已废弃，不得回退"
    )
    assert "all(" in src and "matchable_pairs" in src, "判定必须是全部可比字段逐字相等"

    import app.services.import_template_service as its_mod

    assert "row_idx <= 6" in inspect.getsource(its_mod), "示例行判定仍仅作用于前 6 行"
    # 富模板 AJE/RJE 各 2 行示例 → 除 TEMPLATE_COLUMNS 单示例外另有 3 行须登记，
    # 否则收紧判定会让这 3 行示例变成脏数据（Property 5）。
    assert len(its_mod.TEMPLATE_EXAMPLE_ROWS[ImportType.adjustments]) == 3


def test_a_builtin_example_rows_are_skipped():
    """模板内置示例行必须被识别为示例并跳过 —— 这条基线在 Task 1.3 收紧后仍须为真（Property 5）。"""
    assert _is_example_row(_BUILTIN_EXAMPLE_ROW_1, _ADJ_COLUMNS) is True
    assert _is_example_row(_BUILTIN_EXAMPLE_ROW_2, _ADJ_COLUMNS) is True


def test_a_partially_matching_real_data_row_is_imported_as_data():
    """Task 1.3 新契约：部分字段与示例相同的真实数据行必须按数据行导入，不再被静默丢弃。

    场景：审计师在示例行位置覆盖填写第一笔真实分录（编号仍 AJE-001、类型 AJE、
    科目恰好也是「应收账款」，但摘要与金额都是真实值）→ 3/5 可比字段相同。
    （原断言锁定的是「过半即示例 → 静默丢弃」的旧 bug 行为，Task 1.3 已翻转 basis。）
    """
    real_data_row = [
        "AJE-001", "AJE", "计提本期坏账", "", "应收账款", "", "", "", "0", "888888",
    ]
    assert _is_example_row(real_data_row, _ADJ_COLUMNS) is False


def test_a_fully_different_row_is_data_row():
    """与示例毫无相同可比字段的行 → 数据行（当前与收紧后均为 False）。"""
    row = ["RJE-007", "RJE", "长投重分类", "", "长期股权投资", "", "", "", "50000", "0"]
    assert _is_example_row(row, _ADJ_COLUMNS) is False


def test_a_empty_row_values_is_not_example():
    assert _is_example_row([], _ADJ_COLUMNS) is False


# ═══════════════════════════════════════════════════════════════════════════
# B. 中央 mode=append 导入 `_import_adjustments`
# ═══════════════════════════════════════════════════════════════════════════

_CHART_ROWS = [
    # (account_code, account_name, source)
    ("1122", "应收账款", AccountSource.standard),
    ("1231", "坏账准备", AccountSource.standard),
    ("6602", "管理费用", AccountSource.standard),
    ("112201", "应收账款-客户A", AccountSource.client),
]
_MAPPING_ROWS = [("112201", "1122")]


def _make_db() -> _FakeDB:
    """`_import_adjustments` 恰好按序执行两次查询：AccountChart → AccountMapping。"""
    return _FakeDB([_CHART_ROWS, _MAPPING_ROWS])


class _RecordingSvc:
    """AdjustmentService 替身：记录 create_entry 调用，可按编号注入失败。"""

    calls: list[tuple] = []
    fail_for: set[str] = set()

    def __init__(self, _db):
        pass

    async def create_entry(self, project_id, data, user_id):
        desc = data.description or ""
        if any(k and k in desc for k in _RecordingSvc.fail_for):
            raise ValueError("借贷不平衡")
        _RecordingSvc.calls.append((project_id, data, user_id))
        return SimpleNamespace(id=uuid4())


@pytest.fixture()
def recording_service(monkeypatch):
    import app.services.adjustment_service as adj_mod

    _RecordingSvc.calls = []
    _RecordingSvc.fail_for = set()
    monkeypatch.setattr(adj_mod, "AdjustmentService", _RecordingSvc)
    return _RecordingSvc


async def _run_import(rows: list[dict]) -> dict:
    from app.routers.import_templates import _import_adjustments

    return await _import_adjustments(
        rows, uuid4(), 2025, SimpleNamespace(id=uuid4()), _make_db(),
    )


def test_b_import_adjustments_receives_mode_from_dispatch():
    """Task 1.2 新契约：`_import_adjustments` 有 keyword-only `mode`（默认 append），
    且 `_dispatch_import` 的 adjustments 分支真把 mode 传下去。

    （原断言锁定的是「无 mode 形参 / 分发不传 mode」的旧行为，Task 1.2 已翻转 basis。）
    """
    from app.routers import import_templates as mod

    sig = inspect.signature(mod._import_adjustments)
    params = list(sig.parameters)
    assert params == ["rows", "project_id", "year", "user", "db", "mode"]
    mode_param = sig.parameters["mode"]
    assert mode_param.kind is inspect.Parameter.KEYWORD_ONLY, (
        "mode 设为 keyword-only，既保留既有 5 个位置参数契约，又避免与 user/db 位置冲突"
    )
    assert mode_param.default == "append", "缺省必须是 append（保持既有行为）"

    src = inspect.getsource(mod._dispatch_import)
    assert "_import_adjustments(rows, _require_project(), y, user, db, mode=mode)" in src, (
        "adjustments 分支必须把 mode 传入实现层，否则覆盖模式对调整分录无效"
    )


async def test_b_group_by_adjustment_no_merges_rows(recording_service):
    """同一「编号」多行合并为一个分录组；imported 按分录组计数。"""
    rows = [
        {"编号": "AJE-010", "类型": "AJE", "摘要": "补提坏账", "科目名称": "应收账款",
         "科目编码": "1122", "借方金额": "100", "贷方金额": ""},
        {"编号": "AJE-010", "类型": "AJE", "摘要": "补提坏账", "科目名称": "坏账准备",
         "科目编码": "1231", "借方金额": "", "贷方金额": "100"},
    ]
    result = await _run_import(rows)

    assert result == {"imported": 1, "skipped": 0, "failed": 0, "failed_rows": []}
    assert len(recording_service.calls) == 1
    _, data, _ = recording_service.calls[0]
    assert data.adjustment_type.value == "aje"
    assert data.year == 2025
    assert data.description == "补提坏账"
    assert [li.standard_account_code for li in data.line_items] == ["1122", "1231"]
    assert [li.debit_amount for li in data.line_items] == [Decimal("100"), Decimal("0")]
    assert [li.credit_amount for li in data.line_items] == [Decimal("0"), Decimal("100")]


async def test_b_append_mode_reports_no_skips(recording_service):
    """Task 1.2 新契约：`append` 模式不做 by-key 覆盖判定 → skipped 恒 0，
    且返回结构不带 `skipped_rows` 键（逐字节与改动前一致，Property 2）。

    覆盖跳过只在 `mode=overwrite` 下产生（见 test_adjustment_import_mode.py）。
    （原断言"skipped 恒 0"是对全部模式的描述，Task 1.2 后收窄为仅 append 成立。）
    """
    result = await _run_import([
        {"编号": "AJE-1", "类型": "AJE", "摘要": "x", "科目名称": "应收账款", "借方金额": "1"},
    ])
    assert result["skipped"] == 0
    assert "skipped_rows" not in result


async def test_b_legacy_alias_field_names_supported(recording_service):
    """旧字段名兼容：分录编号 / 调整类型 / 借方科目名称+借方科目代码 / 贷方科目名称+贷方科目代码。"""
    rows = [
        {"分录编号": "RJE-001", "调整类型": "RJE", "摘要": "重分类",
         "借方科目名称": "管理费用", "借方科目代码": "6602", "借方金额": "500"},
        {"分录编号": "RJE-001", "调整类型": "RJE",
         "贷方科目名称": "应收账款", "贷方科目代码": "1122", "贷方金额": "500"},
    ]
    result = await _run_import(rows)

    assert result["imported"] == 1 and result["failed"] == 0
    _, data, _ = recording_service.calls[0]
    assert data.adjustment_type.value == "rje"
    assert [li.standard_account_code for li in data.line_items] == ["6602", "1122"]
    assert [li.account_name for li in data.line_items] == ["管理费用", "应收账款"]


async def test_b_detail_account_code_kept_only_when_differs_from_level1(recording_service):
    """detail_account_code：二级码 ≠ 归一后一级码 → 保留；等于一级码 → None。"""
    rows = [
        # 二级客户码 112201 → 归一到一级 1122，明细码保留
        {"编号": "AJE-020", "类型": "AJE", "摘要": "明细码",
         "科目编码": "112201", "科目名称": "应收账款-客户A", "借方金额": "300"},
        # 一级码 1231 == 归一结果 → 明细码 None（不冗余）
        {"编号": "AJE-020", "类型": "AJE",
         "科目编码": "1231", "科目名称": "坏账准备", "贷方金额": "300"},
    ]
    await _run_import(rows)

    _, data, _ = recording_service.calls[0]
    assert [li.standard_account_code for li in data.line_items] == ["1122", "1231"]
    assert [li.detail_account_code for li in data.line_items] == ["112201", None]


async def test_b_rows_without_no_become_independent_groups(recording_service):
    """未填编号的行各自独立成组（__auto_n），imported 按组数计。"""
    rows = [
        {"类型": "AJE", "摘要": "无编号1", "科目名称": "应收账款", "借方金额": "10"},
        {"类型": "AJE", "摘要": "无编号2", "科目名称": "坏账准备", "贷方金额": "10"},
    ]
    result = await _run_import(rows)
    assert result["imported"] == 2
    assert len(recording_service.calls) == 2


async def test_b_fully_empty_rows_are_ignored(recording_service):
    """完全空行（无科目且借贷均 0）被跳过且不计入任何计数。"""
    rows = [
        {"编号": "", "类型": "", "摘要": "", "科目名称": "", "科目编码": "",
         "借方金额": "", "贷方金额": ""},
        {"编号": "AJE-030", "类型": "AJE", "摘要": "真实行", "科目名称": "应收账款",
         "借方金额": "1"},
    ]
    result = await _run_import(rows)
    assert result == {"imported": 1, "skipped": 0, "failed": 0, "failed_rows": []}


async def test_b_failed_group_reports_friendly_error(recording_service):
    """create_entry 抛错 → failed 计数 + failed_rows 含 adj_no / rows_in_group / error。"""
    recording_service.fail_for = {"会失败"}
    rows = [
        {"编号": "AJE-040", "类型": "AJE", "摘要": "会失败", "科目名称": "应收账款",
         "借方金额": "10"},
    ]
    result = await _run_import(rows)

    assert result["imported"] == 0
    assert result["failed"] == 1
    assert len(result["failed_rows"]) == 1
    fr = result["failed_rows"][0]
    assert fr["adj_no"] == "AJE-040"
    assert fr["rows_in_group"] == [1]
    assert "借贷不平衡" in fr["error"]


async def test_b_unresolvable_account_yields_friendly_hint(recording_service):
    """科目在项目库找不到时，失败信息给出「在项目科目库中找不到」的可读提示。"""
    recording_service.fail_for = {"未知科目"}
    rows = [
        {"编号": "AJE-050", "类型": "AJE", "摘要": "未知科目",
         "科目名称": "不存在的科目", "科目编码": "9999", "借方金额": "10"},
    ]
    result = await _run_import(rows)
    err = result["failed_rows"][0]["error"]
    assert "在项目科目库中找不到" in err
    assert "项目科目库" in err


# ═══════════════════════════════════════════════════════════════════════════
# C. 富模板 export-template 列集合 + export-summary 列头
# ═══════════════════════════════════════════════════════════════════════════

_RICH_TEMPLATE_SHEETS = ["关注事项", "AJE模板", "RJE模板", "项目科目库"]
_RICH_ENTRY_HEADERS = [
    "编号", "类型", "摘要",
    "二级科目编码", "二级科目名称",
    "一级科目编码", "一级科目名称", "报表项目",
    "借方金额", "贷方金额",
]
_RICH_LIB_HEADERS = [
    "二级科目编码", "二级科目名称",
    "一级科目编码", "一级科目名称",
    "报表项目",
    "当前余额", "已有调整数",
]


async def _build_rich_template(monkeypatch) -> openpyxl.Workbook:
    """直接调用 export-template 端点函数（fake db + 跳过自愈加载标准科目）。"""
    from app.routers import adjustments as adj_router
    from app.services import account_chart_service

    async def _noop(*_a, **_kw):
        return None

    monkeypatch.setattr(account_chart_service, "load_standard_template", _noop)

    db = _FakeDB([
        [SimpleNamespace(template_type="soe")],          # Project
        _CHART_ROWS,                                      # AccountChart
        [("1122", "BS-005", "应收账款")],                  # ReportLineMapping
        _MAPPING_ROWS,                                    # AccountMapping
        [("1122", 1000, 900, 50, 0)],                     # TrialBalance
    ])

    resp = await adj_router.export_adjustment_template(
        project_id=uuid4(), year=2025, template_type=None, db=db, current_user=None,
    )
    chunks = []
    async for chunk in resp.body_iterator:
        chunks.append(chunk if isinstance(chunk, bytes) else str(chunk).encode())
    return openpyxl.load_workbook(BytesIO(b"".join(chunks)))


async def test_c_rich_template_sheet_names(monkeypatch):
    """富模板 sheet 名集合固定为 4 张（后续 spec 明令不得改富模板结构）。"""
    wb = await _build_rich_template(monkeypatch)
    assert wb.sheetnames == _RICH_TEMPLATE_SHEETS


async def test_c_rich_template_entry_sheet_headers(monkeypatch):
    """AJE模板 / RJE模板 列头 = 9 列联动结构（10 个列头），逐字不变。"""
    wb = await _build_rich_template(monkeypatch)
    for sheet in ("AJE模板", "RJE模板"):
        ws = wb[sheet]
        headers = [ws.cell(row=1, column=c).value for c in range(1, len(_RICH_ENTRY_HEADERS) + 1)]
        assert headers == _RICH_ENTRY_HEADERS, f"{sheet} 列头漂移"


async def test_c_rich_template_account_library_headers(monkeypatch):
    """项目科目库 sheet 7 列（E 列下拉 + D/F/G/H 联动公式取值源）。"""
    wb = await _build_rich_template(monkeypatch)
    ws = wb["项目科目库"]
    headers = [ws.cell(row=1, column=c).value for c in range(1, len(_RICH_LIB_HEADERS) + 1)]
    assert headers == _RICH_LIB_HEADERS


def test_c_export_summary_headers_include_type_column():
    """Task 1.4 新契约：export-summary 的 `_write_adj_sheet` 列头 8 列，「类型」插在「编号」之后。

    → 原先类型仅隐含在 sheet 名（AJE审计调整 / RJE重分类），导出的汇总改完无法直接导回；
      现列头自带「类型」，使汇总产物可被中央导入直接接受（Property 6 / Property 8）。
    （原断言锁定的是「7 列且无类型」的旧行为，Task 1.4 已翻转 basis。）
    """
    from app.routers import adjustments as adj_router

    src = inspect.getsource(adj_router._write_adj_sheet)
    assert (
        'headers = ["编号", "类型", "摘要", "科目编码", "科目名称", '
        '"借方金额", "贷方金额", "来源"]'
    ) in src
    assert 'headers = ["编号", "摘要"' not in src, "旧 7 列列头不得回退"

    # sheet 名仍保留（缺「类型」列的旧文件靠它兜底推断，见 Property 7）
    export_src = inspect.getsource(adj_router.export_adjustment_summary)
    assert 'ws_aje.title = "AJE审计调整"' in export_src
    assert 'wb.create_sheet("RJE重分类")' in export_src


# ═══════════════════════════════════════════════════════════════════════════
# D. 7 张「无漂移」调整 sheet 的 Sheet_Spec 基线（Property 13）
# ═══════════════════════════════════════════════════════════════════════════

_FACTORY_DEFAULT_STORAGE_FIELD = "conclusion"

# sheet → (模块路径, _SPECS 属性名)
_NO_DRIFT_SHEETS: dict[str, tuple[str, str]] = {
    "K1-4": ("app.routers.wp_render_strategies._k1_import_export", "_K1_SPECS"),
    "K6-3": ("app.routers.wp_render_strategies._k6_import_export", "_K6_SPECS"),
    "K9-3": ("app.routers.wp_render_strategies._k9_import_export", "_K9_SPECS"),
    "K11-3": ("app.routers.wp_render_strategies._k11_import_export", "_K11_SPECS"),
    "K13-3": ("app.routers.wp_render_strategies._k13_import_export", "_K13_SPECS"),
    "I5-3": ("app.routers.wp_render_strategies._i5_import_export", "_I5_SPECS"),
    "I6-3": ("app.routers.wp_render_strategies._i6_import_export", "_I6_SPECS"),
}

_EXPECTED_ITEM_ID = {
    "K1-4": "K1-4-adj-entries",
    "K6-3": "K6-3-adj-entries",
    "K9-3": "K9-3-adj-entries",
    "K11-3": "K11-3-adj-entries",
    "K13-3": "K13-3-adj-entries",
    "I5-3": "I5-3-rows",
    "I6-3": "I6-3-rows",
}

# 有效 storage_field = spec 显式声明 → 否则 router 调用处显式传参 → 否则工厂默认 conclusion
#
# 🔴 characterization 发现（与 design 施工表「已核实无漂移」不一致，如实记录当前值）：
#    K1-4 / K6-3 / I5-3 / I6-3 的有效 storage_field 是 **conclusion**，
#    而这 4 张的前端（useK1Adjustment / K6TabAdjustment / useI5Adjustment / useI6Adjustment）
#    读写的是 **remark** → 属第三种 Orphan_Key 形态（item_id 对但写错列）。
#    本用例只锁定当前真实值；是否纳入 Wave 1 对齐由 spec 决定，届时同步更新本断言。
#
#    ✅ Task 2.7 已修（basis 改变，非放宽断言）：这 4 张各加 `"storage_field": "remark"`，
#       与前端读写列一致；`item_id` / `headers` / `field_keys` 一律未动（下方三条断言不变）。
_EXPECTED_EFFECTIVE_STORAGE_FIELD = {
    "K1-4": "remark",       # Task 2.7 补声明（前端 useK1Adjustment 读写 remark）
    "K6-3": "remark",       # Task 2.7 补声明（前端 K6TabAdjustment 读写 remark）
    "K9-3": "remark",       # spec 显式声明
    "K11-3": "remark",      # router 级 storage_field="remark"
    "K13-3": "remark",      # spec 显式声明
    "I5-3": "remark",       # Task 2.7 补声明（前端 useI5Adjustment 读写 remark）
    "I6-3": "remark",       # Task 2.7 补声明（前端 useI6Adjustment 读写 remark）
}

_EXPECTED_HEADERS = {
    "K1-4": ["调整事项说明", "类别", "报表项目", "科目代码", "科目名称", "附注项目",
             "借方调整金额", "贷方调整金额", "索引", "备注"],
    "K6-3": ["序号", "分录类型", "摘要", "科目代码", "科目名称", "借方金额", "贷方金额",
             "编制人", "备注"],
    "K9-3": ["调整事项说明", "类别", "报表项目", "科目名称", "附注项目", "摘要",
             "借方调整金额", "贷方调整金额", "索引", "备注"],
    "K11-3": ["调整事项说明", "类别", "报表项目", "科目名称", "附注项目", "……",
              "借方调整金额", "贷方调整金额", "索引", "备注"],
    "K13-3": ["类型", "调整事项说明", "类别", "报表项目", "科目名称", "附注项目",
              "借方调整金额", "贷方调整金额", "索引", "备注"],
    "I5-3": ["调整事项说明", "类别", "报表项目", "科目名称", "附注项目", "摘要",
             "借方调整金额", "贷方调整金额", "索引", "备注"],
    "I6-3": ["调整事项说明", "类别", "报表项目", "科目名称", "附注项目", "摘要",
             "借方调整金额", "贷方调整金额", "索引", "备注"],
}

_EXPECTED_FIELD_KEYS = {
    "K1-4": ["description", "category", "reportItem", "accountCode", "accountName",
             "noteItem", "debitAmount", "creditAmount", "indexRef", "remark"],
    "K6-3": ["seq", "entryType", "summary", "accountCode", "accountName",
             "debitAmount", "creditAmount", "preparedBy", "remark"],
    "K9-3": ["description", "category", "reportItem", "accountName", "noteItem",
             "summary", "debitAmount", "creditAmount", "indexRef", "remark"],
    "K11-3": ["description", "category", "reportItem", "accountName", "noteItem",
              "summary", "debitAmount", "creditAmount", "indexRef", "remark"],
    "K13-3": ["type", "description", "category", "reportItem", "accountName",
              "noteItem", "debitAmount", "creditAmount", "refIndex", "remark"],
    "I5-3": ["description", "category", "reportItem", "accountName", "noteItem",
             "summary", "debitAmount", "creditAmount", "indexRef", "remark"],
    "I6-3": ["description", "category", "reportItem", "accountName", "noteItem",
             "summary", "debitAmount", "creditAmount", "indexRef", "remark"],
}


def _load_spec(sheet: str) -> tuple[dict, object]:
    import importlib

    mod_path, specs_attr = _NO_DRIFT_SHEETS[sheet]
    mod = importlib.import_module(mod_path)
    specs = getattr(mod, specs_attr)
    assert sheet in specs, f"{specs_attr} 缺少 {sheet} 条目"
    return specs[sheet], mod


def _router_call_args(mod) -> str:
    """截取 `create_cycle_import_export_router(...)` 调用的实参文本（括号配平）。"""
    src = inspect.getsource(mod)
    marker = "create_cycle_import_export_router("
    idx = src.index(marker)
    i = idx + len(marker)
    depth = 1
    while i < len(src) and depth:
        if src[i] == "(":
            depth += 1
        elif src[i] == ")":
            depth -= 1
        i += 1
    return src[idx + len(marker):i - 1]


def _effective_storage_field(sheet: str) -> str:
    spec, mod = _load_spec(sheet)
    if "storage_field" in spec:
        return spec["storage_field"]
    m = re.search(r'storage_field\s*=\s*"([^"]*)"', _router_call_args(mod))
    if m:
        return m.group(1)
    return _FACTORY_DEFAULT_STORAGE_FIELD


def test_d_factory_storage_field_default_is_conclusion():
    """工厂 `create_cycle_import_export_router` 的 storage_field 默认值 = conclusion。

    这是「只改 item_id 仍读不到」的隐蔽 Orphan_Key 形态的根源，Wave 1 依赖此前提。
    """
    from app.routers.wp_render_strategies._cycle_import_export_common import (
        create_cycle_import_export_router,
    )

    sig = inspect.signature(create_cycle_import_export_router)
    assert sig.parameters["storage_field"].default == _FACTORY_DEFAULT_STORAGE_FIELD


@pytest.mark.parametrize("sheet", sorted(_NO_DRIFT_SHEETS))
def test_d_no_drift_sheet_item_id(sheet):
    spec, _ = _load_spec(sheet)
    assert spec["item_id"] == _EXPECTED_ITEM_ID[sheet]


@pytest.mark.parametrize("sheet", sorted(_NO_DRIFT_SHEETS))
def test_d_no_drift_sheet_effective_storage_field(sheet):
    """有效 storage_field（spec 显式 → router 显式 → 工厂默认）逐字锁定。"""
    assert _effective_storage_field(sheet) == _EXPECTED_EFFECTIVE_STORAGE_FIELD[sheet]


@pytest.mark.parametrize("sheet", sorted(_NO_DRIFT_SHEETS))
def test_d_no_drift_sheet_headers(sheet):
    spec, _ = _load_spec(sheet)
    assert list(spec["headers"]) == _EXPECTED_HEADERS[sheet]


@pytest.mark.parametrize("sheet", sorted(_NO_DRIFT_SHEETS))
def test_d_no_drift_sheet_field_keys(sheet):
    spec, _ = _load_spec(sheet)
    assert list(spec["field_keys"]) == _EXPECTED_FIELD_KEYS[sheet]


@pytest.mark.parametrize("sheet", sorted(_NO_DRIFT_SHEETS))
def test_d_no_drift_sheet_headers_keys_same_length(sheet):
    spec, _ = _load_spec(sheet)
    assert len(spec["headers"]) == len(spec["field_keys"]), f"{sheet} headers/field_keys 长度不一致"

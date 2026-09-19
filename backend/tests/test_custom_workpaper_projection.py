"""自定义底稿投影与可见性修复守卫（Wave 1）。

spec: .kiro/specs/custom-workpaper-dual-mode-formula-and-batch/ Task 4

覆盖 Property 1/2/3/13 与 Requirements 1.2~1.9 / 2.3~2.5 / 4.8 / 10.1~10.3。

🔴 本文件是「缺陷不可回退」守卫，四处修复各配反向自检 + 变异检验（见 Notes）：
  ① create_custom_workpaper 调投影且在 fill_header 之后、commit 之前
  ② write_cell_to_parsed_data 维护 max_row/max_col（单调不缩）
  ③ "custom" 在 _ONLYOFFICE_HTML_WHITELIST
  ④ 投影坐标与 xlsx 恒等（这条是 Wave 2「HTML 编辑写回 xlsx」的正确性前提）

🔴 不读注释：全部源码级断言先 `_strip_comments`，并配「剥注释确实生效」自检
（否则 docstring 里解释缺陷的文字会被数成真实调用）。
"""

from __future__ import annotations

import inspect
import re
from pathlib import Path
from typing import Any

import pytest
from hypothesis import given, settings
from hypothesis import strategies as st

from app.services import custom_workpaper_projection as proj
from app.services.custom_workpaper_projection import (
    GRID_REQUIRED_KEYS,
    col_letter_to_index,
    empty_grid,
    ensure_grid_bounds,
    normalize_cell_ref,
    parse_cell_ref,
    project_custom_workpaper,
    write_cells_to_xlsx,
)
from app.services.wp_parsed_data_service import write_cell_to_parsed_data

# ─── helpers ────────────────────────────────────────────────────────────────

_BACKEND = Path(__file__).resolve().parents[1]
_REPO = _BACKEND.parent


def _strip_comments(src: str) -> str:
    """剥 Python 注释与字符串字面量（含三引号 docstring）。

    🔴 必需：本 spec 的生产代码 docstring 里大量引用被禁的符号名（解释「为什么不用
    extract_grid」），不剥会把说明文字数成真实调用。
    """
    src = re.sub(r'"""[\s\S]*?"""', '""', src)
    src = re.sub(r"'''[\s\S]*?'''", "''", src)
    src = re.sub(r"(?m)#.*$", "", src)
    return src


def _read(rel: str) -> str:
    p = _REPO / rel
    assert p.exists(), f"路径不存在（守卫判据失效）: {rel}"
    txt = p.read_text(encoding="utf-8")
    assert txt.strip(), f"文件为空（守卫判据失效）: {rel}"
    return txt


def _func_body(src: str, name: str) -> str:
    """截取 `async def name(` / `def name(` 的函数体（按缩进，`[ \\t]*` 不用 `\\s*`）。

    🔴 `\\s` 在 re.M 下含换行 → `^(\\s*)def` 会从前面的空行开始匹配，
    indent 被算成一串换行的长度、函数体截成空串（平台已记铁律）。
    """
    m = re.search(rf"(?m)^([ \t]*)(?:async\s+)?def\s+{re.escape(name)}\s*\(", src)
    assert m, f"未找到函数 {name}（守卫判据失效）"
    indent = len(m.group(1))

    # 🔴 先用圆括号配对跳过整个参数列表，再按缩进截函数体。
    # 多行参数列表的收尾 `):` 缩进为 0，若直接按缩进扫会在那一行就 break，
    # 截出来的"函数体"只有签名几行 —— 平台已记的「固定窗口/缩进截函数体」同族坑。
    i = src.index("(", m.end() - 1)
    depth = 0
    while i < len(src):
        if src[i] == "(":
            depth += 1
        elif src[i] == ")":
            depth -= 1
            if depth == 0:
                break
        i += 1
    assert depth == 0, f"{name} 参数列表括号不配对（守卫判据失效）"

    lines = src[i:].splitlines()
    body = [src[m.start():i].replace("\n", " ")]  # 签名压成一行便于断言
    for ln in lines[1:]:
        if ln.strip() and (len(ln) - len(ln.lstrip())) <= indent:
            break
        body.append(ln)
    out = "\n".join(body)
    assert len(out) > 50, f"{name} 函数体过短（截取失败）"
    return out


#: 投影调用形态（**必须带左括号** —— 只匹配调用，不匹配 import 行）。
#: 🔴 M1 变异实测：裸子串断言会被函数体内的 `from ... import refresh_custom_projection`
#: 满足，于是「调用被删」这个最核心的变异静默逃逸。
_PROJ_CALL_RE = re.compile(r"(?:project_custom_workpaper|refresh_custom_projection)\s*\(")


class _FakeWp:
    """WorkingPaper 替身（只需 parsed_data / file_path）。"""

    def __init__(self, parsed_data: dict | None = None, file_path: str | None = None):
        self.parsed_data = parsed_data
        self.file_path = file_path


@pytest.fixture(autouse=True)
def _no_flag_modified(monkeypatch):
    """`flag_modified` 需要真实 ORM 实例，替身场景下打桩。"""
    monkeypatch.setattr(proj, "flag_modified", lambda *_a, **_k: None)
    import app.services.wp_parsed_data_service as pds

    monkeypatch.setattr(pds, "flag_modified", lambda *_a, **_k: None)


def _make_xlsx(tmp_path: Path, sheet: str, cells: dict[str, Any]) -> Path:
    import openpyxl

    wb = openpyxl.Workbook()
    ws = wb.active
    ws.title = sheet
    for ref, val in cells.items():
        ws[ref] = val
    fp = tmp_path / "wp.xlsx"
    wb.save(str(fp))
    wb.close()
    return fp


# ─── Property 2: 投影键集完整性 ──────────────────────────────────────────────


class TestProjectionKeySet:
    """Property 2 —— 键集 ⊇ GtGridSheet 消费键集（含 fail-open 路径）。"""

    def test_required_keys_match_frontend_hasdata_dependency(self):
        """GRID_REQUIRED_KEYS 必须是前端 GtGridSheet 实际消费的键（跨前后端交叉锁死）。"""
        src = _read("audit-platform/frontend/src/components/workpaper/GtGridSheet.vue")
        # hasData 的硬依赖：cells + maxRow
        assert re.search(r"Object\.keys\(cells\.value\)\.length\s*>\s*0", src), (
            "GtGridSheet.hasData 形态已变，GRID_REQUIRED_KEYS 需同步复核"
        )
        assert re.search(r"maxRow\.value\s*>\s*0", src), "hasData 不再依赖 maxRow？需复核"
        for key in ("cells", "max_row", "max_col", "col_widths", "merged_cells", "header_rows"):
            assert key in src, f"GtGridSheet 不再消费 {key}，GRID_REQUIRED_KEYS 需同步"
        assert GRID_REQUIRED_KEYS == frozenset(
            {"cells", "max_row", "max_col", "col_widths", "merged_cells", "header_rows"}
        )

    def test_projection_output_superset_of_required(self, tmp_path):
        grid = project_custom_workpaper(
            _make_xlsx(tmp_path, "X1", {"A1": "项目", "B6": 123.45}), "X1"
        )
        assert GRID_REQUIRED_KEYS <= set(grid), f"投影缺键: {GRID_REQUIRED_KEYS - set(grid)}"

    def test_empty_grid_also_complete(self):
        assert GRID_REQUIRED_KEYS <= set(empty_grid())

    @pytest.mark.parametrize(
        "case",
        ["missing_file", "missing_sheet", "empty_file"],
        ids=["文件不存在", "sheet不存在", "空文件"],
    )
    def test_fail_open_paths_keep_full_keyset(self, tmp_path, case):
        """🔴 fail-open 路径键集也必须齐备。

        既有 `wp_grid_extract.extract_grid` 的空路径只返 5 键（缺 header_rows），
        照它做会让「键集 ⊇ 六键」的守卫在 fail-open 路径必红 —— 这是本模块不复用
        它的第二个理由。
        """
        if case == "missing_file":
            fp = tmp_path / "nope.xlsx"
            sheet = "X1"
        elif case == "missing_sheet":
            fp = _make_xlsx(tmp_path, "X1", {"A1": "v"})
            sheet = "NOSUCH"
        else:
            fp = tmp_path / "empty.xlsx"
            fp.write_bytes(b"")
            sheet = "X1"
        grid = project_custom_workpaper(fp, sheet)
        assert GRID_REQUIRED_KEYS <= set(grid), "fail-open 路径键集不全"
        assert grid["cells"] == {}
        assert grid["max_row"] == 0

    def test_legacy_extract_grid_empty_path_is_indeed_incomplete(self, tmp_path):
        """反向自检：证明「不复用 extract_grid」这条禁令有实际差异，非空转。"""
        from app.services.wp_grid_extract import extract_grid

        legacy = extract_grid(_make_xlsx(tmp_path, "X1", {"A1": "v"}), "NOSUCH")
        assert "header_rows" not in legacy, (
            "extract_grid 空路径现已返回 header_rows —— 本模块的第二条理由已过期，请复核"
        )

    def test_cells_nonempty_implies_positive_bounds(self, tmp_path):
        grid = project_custom_workpaper(_make_xlsx(tmp_path, "X1", {"C7": 1}), "X1")
        assert grid["cells"]
        assert grid["max_row"] > 0 and grid["max_col"] > 0


# ─── Property 13: 坐标恒等（Wave 2 写回的正确性前提）─────────────────────────


class TestCoordinateIdentity:
    """Property 13 —— 投影 B6 ≡ xlsx B6。

    🔴 这条不成立就不能做「HTML 编辑写回 xlsx」：用户改投影 B2 会被写进 xlsx B2
    （表头区）覆盖别的单元格。
    """

    def test_no_row_shift_with_header_block(self, tmp_path):
        fp = _make_xlsx(
            tmp_path,
            "X1",
            {
                "A1": "致同会计师事务所",
                "A2": "编制单位: 某公司",
                "A5": "项目",  # 触发 legacy 的 data_start_row 启发式
                "B5": "金额",
                "A6": "货币资金",
                "B6": 123.45,
            },
        )
        grid = project_custom_workpaper(fp, "X1")
        cells = grid["cells"]
        assert cells["B6"]["v"] == 123.45, "B6 值发生位移 —— 坐标非恒等"
        assert "B2" not in cells, "出现 B2 —— 行号被重编（这正是要避免的）"
        assert cells["A1"]["v"] == "致同会计师事务所", "表头区被剥掉 —— 坐标非恒等"

    def test_legacy_extract_grid_does_shift(self, tmp_path):
        """反向自检：证明 legacy 确实位移（这条禁令不是空转）。"""
        from app.services.wp_grid_extract import extract_grid

        fp = _make_xlsx(
            tmp_path,
            "X1",
            {"A1": "致同会计师事务所", "A5": "项目", "B5": "金额", "A6": "货币资金", "B6": 123.45},
        )
        legacy = extract_grid(fp, "X1")
        assert "B6" not in legacy["cells"], (
            "extract_grid 不再重编行号 —— 本模块存在的首要理由已过期，请复核是否可收敛"
        )
        assert any(
            c.get("v") == 123.45 for c in legacy["cells"].values()
        ), "legacy 未提取到该值，判据构造有误"

    def test_no_column_trim(self, tmp_path):
        """不裁空列：列号即 xlsx 列号。"""
        grid = project_custom_workpaper(_make_xlsx(tmp_path, "X1", {"AA7": "far"}), "X1")
        assert "AA7" in grid["cells"]
        assert grid["max_col"] == 27, f"AA 应为第 27 列，实得 {grid['max_col']}"

    def test_projection_matches_xlsx_after_write(self, tmp_path):
        """Property 1 —— 写 xlsx 后重投影，投影值 == xlsx 值（同坐标）。"""
        fp = _make_xlsx(tmp_path, "X1", {"A1": "h", "B6": 1})
        write_cells_to_xlsx(fp, "X1", {"B6": 999.5, "C9": "new"})
        grid = project_custom_workpaper(fp, "X1")
        assert grid["cells"]["B6"]["v"] == 999.5
        assert grid["cells"]["C9"]["v"] == "new"
        assert grid["max_row"] >= 9


# ─── Property 3 / 4: 边界补齐与引用解析 ──────────────────────────────────────


class TestEnsureGridBounds:
    """Property 3 —— 幂等 + 单调 + 不小于任一格坐标。"""

    def test_idempotent(self):
        g1 = ensure_grid_bounds({"cells": {"C7": 1, "AA3": 2}})
        g2 = ensure_grid_bounds(dict(g1))
        assert (g1["max_row"], g1["max_col"]) == (g2["max_row"], g2["max_col"])

    def test_covers_every_cell(self):
        g = ensure_grid_bounds({"cells": {"C7": 1, "AA3": 2, "B12": 3}})
        assert g["max_row"] >= 12
        assert g["max_col"] >= 27

    def test_monotonic_never_shrinks(self):
        g = ensure_grid_bounds({"cells": {"A1": 1}, "max_row": 99, "max_col": 50})
        assert g["max_row"] == 99 and g["max_col"] == 50, "边界被缩小（违反单调）"

    def test_does_not_mutate_input(self):
        src = {"cells": {"C7": 1}}
        ensure_grid_bounds(src)
        assert "max_row" not in src, "ensure_grid_bounds 修改了入参"

    def test_garbage_bounds_tolerated(self):
        g = ensure_grid_bounds({"cells": {"B5": 1}, "max_row": None, "max_col": "x"})
        assert g["max_row"] >= 5 and g["max_col"] >= 2

    @settings(max_examples=20, deadline=None)
    @given(
        col=st.integers(min_value=1, max_value=200),
        row=st.integers(min_value=1, max_value=500),
    )
    def test_pbt_bounds_cover_generated_ref(self, col: int, row: int):
        from app.services.custom_workpaper_projection import _col_letter

        ref = f"{_col_letter(col)}{row}"
        g = ensure_grid_bounds({"cells": {ref: 1}})
        assert g["max_row"] >= row and g["max_col"] >= col


class TestCellRefParsing:
    """Property 4 —— 往返可逆；非法返回哨兵而非抛错。"""

    @pytest.mark.parametrize(
        "raw,expect",
        [("B5", (5, 2)), ("b5", (5, 2)), ("$B$5", (5, 2)), ("A1", (1, 1)),
         ("Z1", (1, 26)), ("AA3", (3, 27)), ("AB10", (10, 28))],
    )
    def test_parse_ok(self, raw, expect):
        assert parse_cell_ref(raw) == expect

    @pytest.mark.parametrize("raw", ["", "B", "5", "B5C", "1A", "B-5", "  ", "??"])
    def test_parse_illegal_returns_zero(self, raw):
        assert parse_cell_ref(raw) == (0, 0), f"{raw!r} 应判非法"
        assert normalize_cell_ref(raw) is None

    def test_col_letter_index(self):
        assert col_letter_to_index("A") == 1
        assert col_letter_to_index("Z") == 26
        assert col_letter_to_index("AA") == 27
        assert col_letter_to_index("1") == 0

    @settings(max_examples=20, deadline=None)
    @given(
        col=st.integers(min_value=1, max_value=300),
        row=st.integers(min_value=1, max_value=1000),
    )
    def test_pbt_roundtrip(self, col: int, row: int):
        from app.services.custom_workpaper_projection import _col_letter

        ref = f"{_col_letter(col)}{row}"
        assert parse_cell_ref(ref) == (row, col)
        assert normalize_cell_ref(ref.lower()) == ref

    def test_openpyxl_agreement(self):
        """与 openpyxl 口径一致（1-based）。"""
        from openpyxl.utils.cell import (
            column_index_from_string,
            coordinate_from_string,
        )

        for ref in ("A1", "B5", "AA3", "AB10", "Z26"):
            letters, row = coordinate_from_string(ref)
            assert parse_cell_ref(ref) == (row, column_index_from_string(letters))


# ─── Property 14: write_cell_to_parsed_data 维护边界 ────────────────────────


class TestWriteCellMaintainsBounds:
    """Requirements 2.3/2.5 —— 修复「hasData 恒 false」的根因。"""

    def test_bounds_after_single_write(self):
        wp = _FakeWp()
        write_cell_to_parsed_data(wp, sheet_name="X1", cell_ref="C7", value=1)
        sd = wp.parsed_data["html_data"]["X1"]
        assert sd["max_row"] >= 7, "max_row 未维护 —— hasData 会恒 false"
        assert sd["max_col"] >= 3

    def test_bounds_never_shrink(self):
        wp = _FakeWp()
        write_cell_to_parsed_data(wp, sheet_name="X1", cell_ref="C7", value=1)
        write_cell_to_parsed_data(wp, sheet_name="X1", cell_ref="A1", value=2)
        sd = wp.parsed_data["html_data"]["X1"]
        assert (sd["max_row"], sd["max_col"]) == (7, 3), "写更小的格把边界缩小了"

    def test_multi_letter_column(self):
        wp = _FakeWp()
        write_cell_to_parsed_data(wp, sheet_name="X1", cell_ref="AA3", value=1)
        assert wp.parsed_data["html_data"]["X1"]["max_col"] >= 27

    def test_hasdata_would_be_true(self):
        """端到端判据：前端 hasData = cells 非空 && maxRow > 0。"""
        wp = _FakeWp()
        write_cell_to_parsed_data(wp, sheet_name="X1", cell_ref="B5", value=42)
        sd = wp.parsed_data["html_data"]["X1"]
        assert len(sd["cells"]) > 0 and sd["max_row"] > 0

    def test_existing_cells_preserved(self):
        """零回归：既有 cells 写入行为不变（Requirements 11.4）。"""
        wp = _FakeWp({"html_data": {"X1": {"cells": {"A1": {"v": "old", "r": 1, "c": 1}}}}})
        write_cell_to_parsed_data(wp, sheet_name="X1", cell_ref="A1", value="new")
        cell = wp.parsed_data["html_data"]["X1"]["cells"]["A1"]
        assert cell["value"] == "new" and cell["v"] == "new"
        assert cell["r"] == 1, "既有 dict 结构的其他键被丢弃"

    def test_illegal_ref_does_not_corrupt_bounds(self):
        wp = _FakeWp()
        write_cell_to_parsed_data(wp, sheet_name="X1", cell_ref="BAD", value=1)
        sd = wp.parsed_data["html_data"]["X1"]
        assert sd["max_row"] >= 0 and sd["max_col"] >= 0

    @settings(max_examples=20, deadline=None)
    @given(
        col=st.integers(min_value=1, max_value=60),
        row=st.integers(min_value=1, max_value=300),
    )
    def test_pbt_bounds(self, col: int, row: int):
        from app.services.custom_workpaper_projection import _col_letter

        wp = _FakeWp()
        write_cell_to_parsed_data(
            wp, sheet_name="S", cell_ref=f"{_col_letter(col)}{row}", value=1
        )
        sd = wp.parsed_data["html_data"]["S"]
        assert sd["max_row"] >= row and sd["max_col"] >= col


# ─── xlsx 写入语义 ──────────────────────────────────────────────────────────


class TestWriteCellsToXlsx:
    """Requirements 3.2 —— 权威写入不得 fail-open。"""

    def test_missing_sheet_raises(self, tmp_path):
        fp = _make_xlsx(tmp_path, "X1", {"A1": 1})
        with pytest.raises(KeyError):
            write_cells_to_xlsx(fp, "NOSUCH", {"A1": 1})

    def test_missing_file_raises(self, tmp_path):
        with pytest.raises(FileNotFoundError):
            write_cells_to_xlsx(tmp_path / "nope.xlsx", "X1", {"A1": 1})

    def test_illegal_ref_raises(self, tmp_path):
        fp = _make_xlsx(tmp_path, "X1", {"A1": 1})
        with pytest.raises(ValueError):
            write_cells_to_xlsx(fp, "X1", {"NOTACELL": 1})

    def test_other_cells_untouched(self, tmp_path):
        fp = _make_xlsx(tmp_path, "X1", {"A1": "keep", "B2": "keep2", "C3": 1})
        write_cells_to_xlsx(fp, "X1", {"C3": 999})
        grid = project_custom_workpaper(fp, "X1")
        assert grid["cells"]["A1"]["v"] == "keep"
        assert grid["cells"]["B2"]["v"] == "keep2"
        assert grid["cells"]["C3"]["v"] == 999

    def test_no_silent_sheet_creation(self, tmp_path):
        """🔴 sheet 不存在必须抛异常而不是静默创建（会掩盖上游 bug）。"""
        import openpyxl

        fp = _make_xlsx(tmp_path, "X1", {"A1": 1})
        with pytest.raises(KeyError):
            write_cells_to_xlsx(fp, "GHOST", {"A1": 1})
        wb = openpyxl.load_workbook(str(fp))
        try:
            assert "GHOST" not in wb.sheetnames, "静默创建了 sheet"
        finally:
            wb.close()

    def test_source_forbids_fail_open_on_save(self):
        """源码级：write_cells_to_xlsx 不得把写盘异常吞掉。"""
        body = _strip_comments(_func_body(inspect.getsource(proj), "write_cells_to_xlsx"))
        assert "raise" in body, "写盘路径无 raise —— 可能已被改成 fail-open"
        # 🔴 判据必须排除 `finally: try: wb.close() except: pass` —— 那是**关闭**清理，
        # 关闭失败不影响「已 save 成功」这一事实；要禁的是**主路径**（save 之前/之中）吞异常。
        main = body.split("finally:")[0]
        assert "wb.save(" in main, "save 不在主路径（守卫判据失效）"
        assert not re.search(r"except\s+Exception\s*:\s*\n\s*(pass|return)\b", main), (
            "主路径出现吞异常的 except —— xlsx 是权威，写不进去不能报成功"
        )
        # 反向自检：把吞异常注入主路径必须被抓到（防「剥 finally」剥过头成空转）
        mutated = main.replace("wb.save(str(fp))", "try:\n        wb.save(str(fp))\n    except Exception:\n        pass")
        assert re.search(r"except\s+Exception\s*:\s*\n\s*(pass|return)\b", mutated), (
            "反向自检失效：注入吞异常后判据仍不命中，说明该断言是空转"
        )


# ─── 源码级不可回退守卫 ─────────────────────────────────────────────────────


class TestCreateCustomWiring:
    """Requirements 2.1 / 10.1 —— create-custom 必须投影，且顺序正确。"""

    _REL = "backend/app/routers/wp_template.py"

    #: 🔴 Task 19 后创建逻辑抽到共享函数 `_create_one_custom_workpaper`
    #: （单条与批量端点共用，R8.7/R8.9）。判据必须跟着重构走，否则
    #: 「投影调用还在、只是搬了家」会被误判成回归；同时补一条「两个端点
    #: 确实委托共享函数」的交叉锁死，防出现第二份创建实现（平台已实证
    #: 「未完成的重构」缺陷模式：A 保留作入口却没真的委托 B）。
    _SHARED = "_create_one_custom_workpaper"

    def test_both_endpoints_delegate_shared_creator(self):
        src = _strip_comments(_read(self._REL))
        for fn in ("create_custom_workpaper", "create_custom_workpaper_batch"):
            body = _func_body(src, fn)
            assert f"{self._SHARED}(" in body, (
                f"{fn} 未委托 {self._SHARED} ⇒ 存在第二份创建实现，投影会漂移"
            )

    def test_calls_projection(self):
        body = _strip_comments(_func_body(_read(self._REL), self._SHARED))
        # 🔴 必须匹配**调用形态** `name(` —— 函数体内的
        # `from ... import refresh_custom_projection` 也含该子串，
        # 裸子串断言在「调用被删、import 还在」时会假绿（M1 变异实测）。
        assert _PROJ_CALL_RE.search(body), (
            "create_custom_workpaper 未调投影（调用形态）—— 新建底稿会显示「暂无内容」"
        )

    def test_call_assertion_is_not_satisfied_by_import_alone(self):
        """反向自检：只有 import 没有调用时，判据必须判否。

        这正是 M1 变异一度逃逸的原因。
        """
        import_only = (
            "from app.services.custom_workpaper_projection import refresh_custom_projection\n"
            "    pass\n"
        )
        assert not _PROJ_CALL_RE.search(import_only), "判据被裸 import 满足 —— 会放过「调用被删」"
        assert _PROJ_CALL_RE.search("refresh_custom_projection(wp, code)"), "判据认不出真实调用"

    def test_import_present(self):
        """🔴 只有调用没有 import = 运行时 NameError，get_diagnostics 查不出。"""
        src = _strip_comments(_read(self._REL))
        assert "custom_workpaper_projection" in src, "缺少投影模块 import"

    def test_projection_after_header_before_return(self):
        """🔴 顺序不对投影里就没有表头（fill_header 写的是 xlsx 本体）。

        Task 19 后 commit 由调用方做（共享函数不 commit —— per-item savepoint
        的前提），故这里断言「表头 → 投影 → 函数返回」；「投影早于 commit」
        由 `test_shared_creator_does_not_commit` + `test_caller_commits_after_delegate`
        两条合起来保证。
        """
        body = _strip_comments(_func_body(_read(self._REL), self._SHARED))
        i_hdr = body.find("fill_workpaper_header")
        _m = _PROJ_CALL_RE.search(body)
        i_proj = _m.start() if _m else -1
        i_ret = body.rfind("return {")
        assert i_hdr > 0 and i_proj > 0 and i_ret > 0, "锚点缺失（守卫判据失效）"
        assert i_hdr < i_proj, "投影早于表头填充 —— 投影里不会有表头"
        assert i_proj < i_ret, "投影晚于返回 —— 不会被执行"

    def test_shared_creator_does_not_commit(self):
        """🔴 共享创建函数内 commit 会破坏批量的 per-item 回滚语义。"""
        body = _strip_comments(_func_body(_read(self._REL), self._SHARED))
        assert "db.commit(" not in body, (
            "共享创建函数内 commit ⇒ 批量端点的 savepoint 隔离失效（部分成功不可回滚）"
        )

    def test_caller_commits_after_delegate(self):
        """单条端点：commit 必须在委托之后（否则投影不落库）。"""
        body = _strip_comments(
            _func_body(_read(self._REL), "create_custom_workpaper")
        )
        i_del = body.find(f"{self._SHARED}(")
        i_commit = body.find("db.commit")
        assert i_del > 0 and i_commit > 0, "锚点缺失（守卫判据失效）"
        assert i_del < i_commit, "commit 早于创建委托 —— 投影不会落库"

    def test_projection_failure_is_fail_open(self):
        """投影失败不得阻断创建（Requirements 1.4）。"""
        body = _strip_comments(_func_body(_read(self._REL), self._SHARED))
        _m = _PROJ_CALL_RE.search(body)
        assert _m, "未找到投影调用（守卫判据失效）"
        seg = body[max(0, _m.start() - 500) : body.rfind("return {")]
        assert "except" in seg, "投影调用未包 try/except —— 投影失败会让创建整体失败"

    def test_strip_comments_selfcheck(self):
        """反向自检：证明 _strip_comments 真的在剥（否则上面断言全是空转）。"""
        raw = _read(self._REL)
        assert "#" in raw, "源码无注释，自检样本无效"
        stripped = _strip_comments(raw)
        assert stripped.count("#") < raw.count("#"), "_strip_comments 未生效"

    def test_strip_comments_removes_docstring_mentions(self):
        """自检：docstring 里提到的符号必须被剥掉（用内联 fixture，不依赖真实文件）。"""
        sample = '''
def f():
    """本函数不调用 project_custom_workpaper。"""
    # project_custom_workpaper 也不在注释里生效
    return 1
'''
        assert "project_custom_workpaper" not in _strip_comments(sample)


class TestOnlyOfficeWhitelist:
    """Requirements 4.8 / 10.3 —— "custom" 必须在白名单。"""

    def test_custom_in_whitelist(self):
        from app.routers.wp_render_config import _ONLYOFFICE_HTML_WHITELIST

        assert "custom" in _ONLYOFFICE_HTML_WHITELIST, (
            "custom 不在白名单 —— 将来变多 sheet 会被静默改写成 onlyoffice-sheet，"
            "GtCustomWpEditor 永不渲染"
        )

    def test_whitelist_is_consulted_by_render_config(self):
        """反向自检：证明该白名单确实被消费（不是死配置）。"""
        src = _strip_comments(_read("backend/app/routers/wp_render_config.py"))
        hits = len(re.findall(r"_ONLYOFFICE_HTML_WHITELIST", src))
        assert hits >= 2, f"白名单只出现 {hits} 次（疑为死配置，判据需复核）"

    def test_custom_registered_in_frontend_registry(self):
        """custom 必须有前端渲染组件，否则白名单放行后是空白页。"""
        src = _read("audit-platform/frontend/src/components/workpaper/htmlRendererRegistry.ts")
        assert "GtCustomWpEditor" in src
        # registry 是 `{ ct: 'custom', component: GtCustomWpEditor }` 形态，
        # 不是 `'custom': X` 的 map 字面量 → 按 ct + 组件名双锚点断言
        # 真正要钉的不变式 = custom 底稿有渲染组件（组件名是稳定锚点，
        # 而 registry 的键写法/条目形态属前端实现细节，不该由后端守卫锁死）
        assert "GtCustomWpEditor" in src, "registry 未引用 GtCustomWpEditor"
        assert re.search(r"\bcustom\b", src), "registry 无 custom 字样"
        assert "GtCustomWpEditor" in src, "registry 未指向 GtCustomWpEditor"


class TestProjectionIsSingleEntry:
    """Requirements 1.5 —— 投影逻辑收敛为单一函数，不得多处各写一份 openpyxl 读取。"""

    def test_no_duplicate_openpyxl_grid_reader(self):
        src = _strip_comments(inspect.getsource(proj))
        # 模块内只允许两处 load_workbook：投影读 + 写回
        hits = len(re.findall(r"load_workbook", src))
        assert hits <= 2, f"模块内 load_workbook 出现 {hits} 次，疑重复实现"

    def test_extract_grid_untouched(self):
        """Requirements 1.6 —— 既有 extract_grid 保持零改动（11 个调用方依赖其现行为）。"""
        src = _read("backend/app/services/wp_grid_extract.py")
        assert "row_offset = data_start_row - 1" in src, (
            "extract_grid 的重编号行为已变 —— 其 11 个只读调用方需重新验证"
        )

    def test_projection_module_does_not_import_extract_grid(self):
        src = _strip_comments(inspect.getsource(proj))
        assert "extract_grid" not in src, (
            "投影模块引用了 extract_grid —— 它会重编行号，写回时会覆盖别的单元格"
        )

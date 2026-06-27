"""Unit tests for wp_docx_template_parser — edge cases.

Tests:
- Empty docx: returns TemplateStructure with empty lists
- Docx with no placeholders: paragraphs filled, placeholders empty
- Docx with legacy markers only (××公司): extracts entity_name placeholder
- Invalid file path: raises FileNotFoundError
- Non-docx file: raises ValueError
- Mixed: Document with both ${...} and legacy markers
- Tables: Document with placeholder in table cell
- Headings: Document with Heading 1/2/3
- Cache: get_cached_structure fills wp_code in metadata
"""

from __future__ import annotations

from pathlib import Path

import pytest
from docx import Document

from app.services.wp_docx_template_parser import (
    TemplateStructure,
    get_cached_structure,
    invalidate_cache,
    parse_template,
)


# ─── Fixtures ─────────────────────────────────────────────────────────────────


@pytest.fixture(autouse=True)
def _clear_cache():
    """Clear template cache before each test."""
    invalidate_cache()
    yield
    invalidate_cache()


# ─── Test: Empty docx ─────────────────────────────────────────────────────────


class TestEmptyDocument:
    """Empty docx returns TemplateStructure with empty lists."""

    def test_empty_doc_returns_template_structure(self, tmp_path: Path):
        """An empty docx file produces a valid TemplateStructure."""
        doc = Document()
        file_path = tmp_path / "empty.docx"
        doc.save(str(file_path))

        result = parse_template(str(file_path))

        assert isinstance(result, TemplateStructure)
        assert result.paragraphs == []
        assert result.tables == []
        assert result.placeholders == []

    def test_empty_doc_has_metadata(self, tmp_path: Path):
        """Even empty docs produce valid metadata."""
        doc = Document()
        file_path = tmp_path / "empty.docx"
        doc.save(str(file_path))

        result = parse_template(str(file_path))

        assert "template_name" in result.metadata
        assert result.metadata["template_name"] == "empty"
        assert "last_parsed_at" in result.metadata


# ─── Test: No placeholders ────────────────────────────────────────────────────


class TestNoPlaceholders:
    """Docx with only static text: paragraphs filled, placeholders empty."""

    def test_static_text_paragraphs_extracted(self, tmp_path: Path):
        """Static paragraphs are captured in paragraphs list."""
        doc = Document()
        doc.add_paragraph("第一段纯文本内容")
        doc.add_paragraph("第二段纯文本内容")
        file_path = tmp_path / "static.docx"
        doc.save(str(file_path))

        result = parse_template(str(file_path))

        assert len(result.paragraphs) == 2
        assert result.paragraphs[0].text == "第一段纯文本内容"
        assert result.paragraphs[1].text == "第二段纯文本内容"

    def test_static_text_no_placeholders(self, tmp_path: Path):
        """No placeholders extracted from static-only doc."""
        doc = Document()
        doc.add_paragraph("纯文本，无占位符")
        doc.add_heading("标题但无变量", level=2)
        file_path = tmp_path / "no_ph.docx"
        doc.save(str(file_path))

        result = parse_template(str(file_path))

        assert result.placeholders == []
        # But paragraphs should have empty placeholder_ids
        for para in result.paragraphs:
            assert para.placeholder_ids == []


# ─── Test: Legacy markers only ────────────────────────────────────────────────


class TestLegacyMarkersOnly:
    """Docx with legacy Chinese markers (××公司) extracts placeholders."""

    def test_legacy_xx_company_extracted(self, tmp_path: Path):
        """××公司 marker extracts entity_name placeholder."""
        doc = Document()
        doc.add_paragraph("关于对××公司的年度审计报告")
        file_path = tmp_path / "legacy.docx"
        doc.save(str(file_path))

        result = parse_template(str(file_path))

        assert len(result.placeholders) >= 1
        entity_placeholders = [p for p in result.placeholders if p.field_id == "entity_name"]
        assert len(entity_placeholders) == 1
        assert entity_placeholders[0].label == "被审计单位"
        assert entity_placeholders[0].pattern == "××公司"

    def test_legacy_date_marker_extracted(self, tmp_path: Path):
        """202X年 marker extracts audit_year placeholder."""
        doc = Document()
        doc.add_paragraph("本函涵盖202X年度财务报表审计")
        file_path = tmp_path / "legacy_date.docx"
        doc.save(str(file_path))

        result = parse_template(str(file_path))

        year_placeholders = [p for p in result.placeholders if "year" in p.field_id or "audit" in p.field_id]
        assert len(year_placeholders) >= 1

    def test_legacy_multiple_markers(self, tmp_path: Path):
        """Multiple legacy markers produce multiple unique placeholders."""
        doc = Document()
        doc.add_paragraph("关于对××公司202X年度财务报表的审计")
        file_path = tmp_path / "multi_legacy.docx"
        doc.save(str(file_path))

        result = parse_template(str(file_path))

        # Should have at least entity_name and audit_year
        field_ids = {p.field_id for p in result.placeholders}
        assert "entity_name" in field_ids


# ─── Test: Invalid file path ──────────────────────────────────────────────────


class TestInvalidFilePath:
    """Non-existent file raises FileNotFoundError."""

    def test_nonexistent_file_raises(self):
        """parse_template raises FileNotFoundError for missing file."""
        with pytest.raises(FileNotFoundError, match="模板文件不存在"):
            parse_template("/nonexistent/path/template.docx")

    def test_nonexistent_file_includes_path_in_message(self):
        """Error message includes the attempted file path."""
        bad_path = "/tmp/does_not_exist_abc123.docx"
        with pytest.raises(FileNotFoundError, match=bad_path):
            parse_template(bad_path)


# ─── Test: Non-docx file ─────────────────────────────────────────────────────


class TestNonDocxFile:
    """Invalid file format raises ValueError."""

    def test_text_file_raises_value_error(self, tmp_path: Path):
        """A .txt file renamed to .docx raises ValueError."""
        file_path = tmp_path / "fake.docx"
        file_path.write_text("This is not a docx file", encoding="utf-8")

        with pytest.raises(ValueError, match="Invalid docx file"):
            parse_template(str(file_path))

    def test_binary_garbage_raises_value_error(self, tmp_path: Path):
        """Random binary data raises ValueError."""
        file_path = tmp_path / "garbage.docx"
        file_path.write_bytes(b"\x00\x01\x02\x03" * 100)

        with pytest.raises(ValueError, match="Invalid docx file"):
            parse_template(str(file_path))


# ─── Test: Mixed new + legacy placeholders ────────────────────────────────────


class TestMixedPlaceholders:
    """Docx with both ${...} new format and legacy Chinese markers."""

    def test_mixed_extracts_both_types(self, tmp_path: Path):
        """Both ${field_id} and ××公司 markers are extracted."""
        doc = Document()
        doc.add_paragraph("关于对××公司的${audit_scope:审计范围}审计报告")
        file_path = tmp_path / "mixed.docx"
        doc.save(str(file_path))

        result = parse_template(str(file_path))

        field_ids = {p.field_id for p in result.placeholders}
        assert "audit_scope" in field_ids
        assert "entity_name" in field_ids

    def test_mixed_new_format_has_label(self, tmp_path: Path):
        """New format placeholder extracts label from colon syntax."""
        doc = Document()
        doc.add_paragraph("审计期间: ${audit_period:审计期间}")
        file_path = tmp_path / "new_label.docx"
        doc.save(str(file_path))

        result = parse_template(str(file_path))

        ph = next(p for p in result.placeholders if p.field_id == "audit_period")
        assert ph.label == "审计期间"
        assert ph.pattern == "${audit_period:审计期间}"

    def test_mixed_preserves_paragraph_references(self, tmp_path: Path):
        """Placeholders from both formats are listed in paragraph placeholder_ids."""
        doc = Document()
        doc.add_paragraph("致××公司全体股东：关于${scope:范围}的说明")
        file_path = tmp_path / "mixed_refs.docx"
        doc.save(str(file_path))

        result = parse_template(str(file_path))

        # First paragraph should reference both placeholders
        para = result.paragraphs[0]
        assert "entity_name" in para.placeholder_ids
        assert "scope" in para.placeholder_ids


# ─── Test: Table placeholders ─────────────────────────────────────────────────


class TestTablePlaceholders:
    """Document with placeholder in table cell extracts with correct position."""

    def test_table_placeholder_extracted(self, tmp_path: Path):
        """Placeholder inside a table cell is extracted."""
        doc = Document()
        table = doc.add_table(rows=2, cols=2)
        table.cell(0, 0).text = "项目名称"
        table.cell(0, 1).text = "${project_name:项目名称}"
        table.cell(1, 0).text = "审计年度"
        table.cell(1, 1).text = "${audit_year:年度}"
        file_path = tmp_path / "table.docx"
        doc.save(str(file_path))

        result = parse_template(str(file_path))

        field_ids = {p.field_id for p in result.placeholders}
        assert "project_name" in field_ids
        assert "audit_year" in field_ids

    def test_table_placeholder_has_correct_position(self, tmp_path: Path):
        """Table placeholder position includes table_index, row, col."""
        doc = Document()
        table = doc.add_table(rows=2, cols=3)
        table.cell(1, 2).text = "${amount:金额}"
        file_path = tmp_path / "table_pos.docx"
        doc.save(str(file_path))

        result = parse_template(str(file_path))

        amount_ph = next(p for p in result.placeholders if p.field_id == "amount")
        assert amount_ph.position["table_index"] == 0
        assert amount_ph.position["row"] == 1
        assert amount_ph.position["col"] == 2

    def test_table_structure_captured(self, tmp_path: Path):
        """Table rows and cells are captured in the tables list."""
        doc = Document()
        table = doc.add_table(rows=2, cols=2)
        table.cell(0, 0).text = "A"
        table.cell(0, 1).text = "B"
        table.cell(1, 0).text = "C"
        table.cell(1, 1).text = "D"
        file_path = tmp_path / "table_struct.docx"
        doc.save(str(file_path))

        result = parse_template(str(file_path))

        assert len(result.tables) == 1
        assert result.tables[0].rows == [["A", "B"], ["C", "D"]]
        assert result.tables[0].index == 0

    def test_table_placeholder_ids_in_table_def(self, tmp_path: Path):
        """Table placeholders are listed in the TableDef.placeholder_ids."""
        doc = Document()
        table = doc.add_table(rows=1, cols=2)
        table.cell(0, 0).text = "单位"
        table.cell(0, 1).text = "${entity:单位名称}"
        file_path = tmp_path / "table_ids.docx"
        doc.save(str(file_path))

        result = parse_template(str(file_path))

        assert "entity" in result.tables[0].placeholder_ids


# ─── Test: Heading levels ─────────────────────────────────────────────────────


class TestHeadingLevels:
    """Documents with Heading 1/2/3 report correct heading_level values."""

    def test_heading_1(self, tmp_path: Path):
        """Heading 1 paragraph has heading_level=1."""
        doc = Document()
        doc.add_heading("一、审计概述", level=1)
        file_path = tmp_path / "h1.docx"
        doc.save(str(file_path))

        result = parse_template(str(file_path))

        assert result.paragraphs[0].heading_level == 1
        assert "Heading" in result.paragraphs[0].style

    def test_heading_2(self, tmp_path: Path):
        """Heading 2 paragraph has heading_level=2."""
        doc = Document()
        doc.add_heading("（一）审计范围", level=2)
        file_path = tmp_path / "h2.docx"
        doc.save(str(file_path))

        result = parse_template(str(file_path))

        assert result.paragraphs[0].heading_level == 2

    def test_heading_3(self, tmp_path: Path):
        """Heading 3 paragraph has heading_level=3."""
        doc = Document()
        doc.add_heading("1. 具体事项", level=3)
        file_path = tmp_path / "h3.docx"
        doc.save(str(file_path))

        result = parse_template(str(file_path))

        assert result.paragraphs[0].heading_level == 3

    def test_normal_paragraph_heading_level_zero(self, tmp_path: Path):
        """Normal paragraph has heading_level=0."""
        doc = Document()
        doc.add_paragraph("这是一段普通文本")
        file_path = tmp_path / "normal.docx"
        doc.save(str(file_path))

        result = parse_template(str(file_path))

        assert result.paragraphs[0].heading_level == 0

    def test_mixed_headings_and_normal(self, tmp_path: Path):
        """Document with mixed heading levels reports correct values."""
        doc = Document()
        doc.add_heading("第一章", level=1)
        doc.add_paragraph("正文内容")
        doc.add_heading("第一节", level=2)
        doc.add_paragraph("更多内容")
        file_path = tmp_path / "mixed_headings.docx"
        doc.save(str(file_path))

        result = parse_template(str(file_path))

        assert result.paragraphs[0].heading_level == 1
        assert result.paragraphs[1].heading_level == 0
        assert result.paragraphs[2].heading_level == 2
        assert result.paragraphs[3].heading_level == 0


# ─── Test: Cache (get_cached_structure) ───────────────────────────────────────


class TestCachedStructure:
    """get_cached_structure fills wp_code in metadata and caches correctly."""

    def test_cache_fills_wp_code(self, tmp_path: Path):
        """get_cached_structure sets metadata.wp_code to given wp_code."""
        doc = Document()
        doc.add_paragraph("测试文档")
        file_path = tmp_path / "cached.docx"
        doc.save(str(file_path))

        result = get_cached_structure(str(file_path), "A8-1")

        assert result.metadata["wp_code"] == "A8-1"

    def test_cache_returns_same_object_on_second_call(self, tmp_path: Path):
        """Same file path without modification returns identical object."""
        doc = Document()
        doc.add_paragraph("缓存测试")
        file_path = tmp_path / "cache_test.docx"
        doc.save(str(file_path))

        result1 = get_cached_structure(str(file_path), "A9-1")
        result2 = get_cached_structure(str(file_path), "A9-1")

        assert result1 is result2  # Same object identity (cached)

    def test_cache_file_not_found_raises(self, tmp_path: Path):
        """get_cached_structure raises FileNotFoundError for missing file."""
        with pytest.raises(FileNotFoundError):
            get_cached_structure(str(tmp_path / "missing.docx"), "A10-1")

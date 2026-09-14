"""G0 投资循环函证 —— 源模板事实固化守卫

spec: g0-confirmation-source-alignment，Requirement 1 / 10（Property 1, 2, 8, 16, 19, 24, 25, 27）

裁决者 = `backend/wp_templates/G/G0 投资循环函证.xlsx`（openpyxl 直读，不连库 → 可进 CI）。

本文件的定位：把逐格精读结论冻结成断言，使后续任何会话都**不必重新精读源模板**，
也不能凭"常识"改动 G0 结构。每类断言配反向自检（Requirement 1.8）。

🔴 三处 tab 名索引号笔误是**源模板事实**，不得"修正"：
  - `函证差异核对表G0-3（证券投资）` —— 底稿目录为 G0-4
  - `函证差异核对表G0-4(非证券投资)` —— 底稿目录为 G0-5
  - `函证程序舞弊风险评价表F0-8`     —— 底稿目录为 G0-8
定位一律用 tab 名，展示用底稿目录索引号（详见 design.md §6）。
"""

from __future__ import annotations

from pathlib import Path

import openpyxl
import pytest

_REPO_ROOT = Path(__file__).resolve().parents[2]
_G0_PATH = _REPO_ROOT / "backend" / "wp_templates" / "G" / "G0 投资循环函证.xlsx"
_D0_PATH = _REPO_ROOT / "backend" / "wp_templates" / "D" / "D0 收入循环函证.xlsx"

# ─── 源模板事实常量（逐格精读，2026-08-03） ──────────────────────────────────

SHEET_NAMES: tuple[str, ...] = (
    "底稿目录",
    "函证程序表G0A",
    "函证结果汇总表G0-1",
    "核实被函证单位信息G0-2",
    "跟函函证过程控制G0-3",
    "函证差异核对表G0-3（证券投资）",
    "函证差异核对表G0-4(非证券投资)",
    "替代程序检查表G0-6",
    "邮件传真回函可靠性验证G0-7",
    "函证程序舞弊风险评价表F0-8",
)

SHEET_SUMMARY = "函证结果汇总表G0-1"
SHEET_ENTITY_VERIFY = "核实被函证单位信息G0-2"
SHEET_FOLLOWUP = "跟函函证过程控制G0-3"
SHEET_DIFF_SECURITIES = "函证差异核对表G0-3（证券投资）"
SHEET_DIFF_NONSECURITIES = "函证差异核对表G0-4(非证券投资)"
SHEET_ALTERNATIVE = "替代程序检查表G0-6"
SHEET_RELIABILITY = "邮件传真回函可靠性验证G0-7"
SHEET_FRAUD = "函证程序舞弊风险评价表F0-8"
SHEET_PROGRAM = "函证程序表G0A"

# 底稿目录（D:F 列，R4:R12）—— 索引号真源
DIRECTORY_ENTRIES: tuple[tuple[str, str], ...] = (
    ("函证程序表", "G0A"),
    ("函证结果汇总表", "G0-1"),
    ("核实被函证单位信息", "G0-2"),
    ("跟函函证过程控制", "G0-3"),
    ("函证差异核对表（证券投资）", "G0-4"),
    ("函证差异核对表(非证券投资)", "G0-5"),
    ("替代程序检查表", "G0-6"),
    ("邮件传真回函可靠性验证", "G0-7"),
    ("函证程序舞弊风险评价表", "G0-8"),
)

# G0-1 上区 R5 段头（合并区起始格 → 文字）
# 🔴 「3、回函金额确认」在 G0-1 落在 U5:W5，而 D0-1/F0-1 均为 S5:W5
#    → S 回函金额 / T 差异 两列无段归属（源模板缺陷 #7，见 TestSourceDefectsExist）
SUMMARY_R5_GROUPS: dict[str, str] = {
    "A5": "序号",
    "B5": "询证函索引号",
    "C5": "发函询证纪要",
    "G5": "1、发函信息",
    "L5": "2、收到回函",
    "U5": "3、回函金额确认",
    "X5": "4、未收到回函的替代程序",
    "AB5": "审计结论",
}

# R5 段头合并区（G0-1 实际值）
SUMMARY_R5_MERGES: tuple[str, ...] = ("C5:F5", "G5:K5", "L5:R5", "U5:W5", "X5:AA5")
# D0-1 / F0-1 的同位合并区（结构同构，作为「正确意图」的旁证）
SUMMARY_R5_MERGES_D0F0: tuple[str, ...] = ("C5:F5", "G5:K5", "L5:R5", "S5:W5", "X5:AA5")

# G0-1 上区 28 个叶子列（A..AB；A/B/AB 在 R5 且纵向合并到 R7，其余在 R6）
SUMMARY_LEAF_COLUMNS: tuple[tuple[str, str], ...] = (
    ("A", "序号"),
    ("B", "询证函索引号"),
    ("C", "选取样本目的"),
    ("D", "被询证单位名称"),
    ("E", "账户/交易"),
    ("F", "账面期末余额"),
    ("G", "函证方式"),
    ("H", "发函日期"),
    ("I", "发函单号"),
    ("J", "收件地址"),
    ("K", "地址核查是否一致"),
    ("L", "是否收到回函（√）"),
    ("M", "回函方式"),
    ("N", "是否相符"),
    ("O", "回函日期"),
    ("P", "回函快递单号"),
    ("Q", "回函地址"),
    ("R", "发函地址与回函地址是否一致"),
    ("S", "回函金额"),
    ("T", "差异"),
    ("U", "可确认金额"),
    ("V", "调节索引（G0-3）"),
    ("W", "其他说明/备注"),
    ("X", "是否采取替代程序（√）"),
    ("Y", "替代后可确认金额"),
    ("Z", "替代后不可确认金额"),
    ("AA", "替代程序索引号"),
    ("AB", "审计结论"),
)

# G0-1 下区四块锚点
SUMMARY_LOWER_BLOCKS: tuple[tuple[str, str], ...] = (
    ("C19", "一、函证情况"),
    ("J19", "二、样本选择"),
    ("S19", "三、审计说明"),
    ("C30", "四、审计结论"),
)

# G0-1「一、函证情况」8 指标（C21:C28 逐字）
SUMMARY_MATRIX_METRICS: tuple[tuple[str, str], ...] = (
    ("C21", "本期（期末）账面金额："),
    ("C22", "抽取样本的发函金额："),
    ("C23", "发函金额占账面金额的比例(%)："),
    ("C24", "回函确认金额："),
    ("C25", "回函可确认金额占发函金额的比例(%)："),
    ("C26", "回函可确认金额占账面金额的比例(%)："),
    ("C27", "替代测试确认金额："),
    ("C28", "回函和替代确认金额占账面金额的比例(%)："),
)

# G0-1 矩阵品种（E20:H20；H20 是 `……` 可扩位）
SUMMARY_MATRIX_CATEGORIES: tuple[tuple[str, str], ...] = (
    ("E20", "交易性金融资产"),
    ("F20", "长期股权投资"),
    ("G20", "债权投资"),
    ("H20", "……"),
)

# G0-1「二、样本选择」6 项标签（J20..J26；J24 无标签、K24 是补充说明）
SUMMARY_SAMPLE_LABELS: tuple[tuple[str, str], ...] = (
    ("J20", "测试总体："),
    ("J21", "特定样本："),
    ("J22", "抽样总体："),
    ("J23", "确定的抽样样本量："),
    ("J25", "抽样方法："),
    ("J26", "抽样过程："),
)

# G0-1「三、审计说明」5 项标题
SUMMARY_AUDIT_NOTE_TITLES: tuple[tuple[str, str], ...] = (
    ("S20", "1、对询证函保持的控制的说明"),
    ("X20", "2、对误差的分析"),
    ("S24", "3、对以传真或电子邮件形式收到的回函的可靠性的考虑（G0-6）"),
    ("S25", "4、针对不符事项的程序"),
    ("S28", "5、针对未回函的替代程序"),
)

# G0-2 三段表头（合并区起始格 → 文字）
ENTITY_VERIFY_R5_GROUPS: dict[str, str] = {
    "A5": "询证函索引号",
    "B5": "被审计单位提供的被函证单位信息及核对（核实被函证单位信息）（注1）",
    "P5": "回函信息情况（回函核对记录）",
    "AB5": "第一次发函结果？（送抵/退回）（注2）",
}

# G0-2 R6 叶子列（B..AA + AF..AK；AB..AE/AL 在 R5 纵向合并）
ENTITY_VERIFY_LEAF_R6: tuple[tuple[str, str], ...] = (
    ("B", "被询证单位全称"),
    ("C", "函证方式"),
    ("D", "收件地址"),
    ("E", "邮编"),
    ("F", "联系人"),
    ("G", "联系电话"),
    ("H", "邮箱/传真"),
    ("I", "企查查地址"),
    ("J", "发函地址与企查查地址是否一致"),
    ("K", "地址核查不一致的说明"),
    ("L", "地址不一致的核实方式"),
    ("M", "核实的支持性文件或其他公开信息索引号"),
    ("N", "不一致说明是否合理"),
    ("O", "备注"),
    ("P", "回函方式"),
    ("Q", "是否为原件"),
    ("R", "是否直接收到回函"),
    ("S", "回函发出地址（含物流信息中的地址）"),
    ("T", "回函寄件人"),
    ("U", "回函电话"),
    ("V", "发函地址与回函地址是否一致"),
    ("W", "发函收件人与回函寄件人是否一致"),
    ("X", "发函收件人电话与回函寄件人电话是否一致"),
    ("Y", "不一致的说明"),
    ("Z", "核实支持性证据索引号"),
    ("AA", "跟函函证控制过程（G0-2）"),
    ("AF", "地址"),
    ("AG", "邮编"),
    ("AH", "联系人"),
    ("AI", "联系电话"),
    ("AJ", "传真"),
    ("AK", "信息是否核查一致"),
)

# G0-1 → G0-2 的 VLOOKUP 列序（源 R8 公式实证）
SUMMARY_VLOOKUP_COL_INDEX: dict[str, int] = {
    "D": 2,   # 被询证单位名称   ← G0-2!B
    "G": 3,   # 函证方式         ← G0-2!C
    "J": 4,   # 收件地址         ← G0-2!D
    "K": 10,  # 地址核查是否一致 ← G0-2!J
    "M": 16,  # 回函方式         ← G0-2!P
    "Q": 19,  # 回函地址         ← G0-2!S
    "R": 22,  # 发函/回函地址一致 ← G0-2!V
}

# G0A 12 条程序（seq, 程序分类 D 列, 底稿索引号 E 列）
PROGRAM_ROWS: tuple[tuple[str, str, str | None], ...] = (
    ("1", "常规★", "G0-1"),
    ("2", "IPO/上市/新三板/重组/舞弊应对", "G0-1"),
    ("3", "常规★", "G0-2"),
    ("4", "常规★", "G0-1/G0-3"),
    ("5", "常规★", "G0-1/G0-2"),
    ("6", "常规★", "G0-1/G0-2/G0-3/G0-4"),
    ("7", "备选", "G0-1/G0-7"),
    ("8", "备选", None),
    ("9", "常规★", "G0-1/G0-2"),
    ("10", "常规★", "G0-6"),
    ("11", "舞弊应对/IPO/上市/新三板/重组", "G0-8"),
    ("12", "常规★", None),
)

# G0-6 区块锚点与两级表头
ALT_HEADER_FIELDS: tuple[tuple[str, str], ...] = (("A5", "会计科目："), ("D5", "投资产品/名称："))

ALT_BLOCK_ANCHORS: tuple[tuple[str, str], ...] = (
    ("A6", "一、样本选取标准与规模"),
    ("A8", "二、检查过程记录"),
    ("A9", "1.检查初始投资协议、公司章程等"),
    ("A15", "2.检查本期发生额"),
    ("A16", "（1）本期借方发生额"),
    ("A24", "（2）本期贷方发生额"),
    ("A32", "3.检查期后是否被出售或赎回"),
    ("A40", "三、审计说明："),
    ("A43", "四、审计结论："),
    ("A46", "编制说明："),
)

ALT_BLOCK1_COLUMNS: tuple[str, ...] = ("被投资单位", "投资比例", "投资金额", "投资条款", "索引号")

# G0-6 两组点选项（纸质底稿的打勾位，平台改为多选）
#
# 🔴 二者**只差最后一项**：`B7` 末项是「全部」、`C15` 末项是「其他」。抄错会让
#    「一、样本选取标准与规模」与「2.检查本期发生额」两处的语义混同。
ALT_TEST_SCOPE_TEXT = "大额（）关联方（）大额交易频繁（）异常（）全部（）"
ALT_OCCURRENCE_SCOPE_TEXT = "大额（）关联方（）大额交易频繁（）异常（）其他（）"
ALT_SCOPE_CELLS: tuple[tuple[str, str], ...] = (
    ("B7", ALT_TEST_SCOPE_TEXT),
    ("C15", ALT_OCCURRENCE_SCOPE_TEXT),
)

# 编制说明三条替代程序要点（正文在 B 列，`A46`/`A47` 只是标题）
ALT_PREPARATION_NOTES: tuple[tuple[str, str], ...] = (
    ("B48", "①检查期初原始投资协议、期后出售或赎回协议；"),
    ("B49", "②检查原始凭证：合同、交易流水、银行回单、支票存根等；"),
    ("B50", "③对回函可能性不高的、余额重大的，发函同时执行替代程序。"),
)

# 🔴 源模板**只有四个带序号的段**，且「二」是检查过程记录（不是余额汇总 —— 平台
#    改造前用源外增强段占了「二」，导致其后全部串位）。
ALT_NUMBERED_SECTIONS: tuple[tuple[str, str], ...] = (
    ("A6", "一、样本选取标准与规模"),
    ("A8", "二、检查过程记录"),
    ("A40", "三、审计说明："),
    ("A43", "四、审计结论："),
)

# 区块② 借/贷两区共用同一两级表头
ALT_BLOCK2_GROUP_ROW: tuple[str | None, ...] = (
    "记账凭证", None, None, None, None,
    "支持性文件1", None, None,
    "支持性文件2", None, None,
    "……", "索引号", "是否异常",
)
ALT_BLOCK2_LEAF_ROW: tuple[str | None, ...] = (
    "日期", "凭证编号", "业务内容", "对方科目", "金额",
    "识别特征", "信息1", "信息2",
    "识别特征", "信息1", "信息2",
    None, None, None,
)

ALT_BLOCK3_GROUP_ROW: tuple[str | None, ...] = (
    "记账凭证", None, None, None, None,
    "投资协议/交易确认单/交割单", None, None,
    "银行回单", None, None,
    "……", "索引号", "是否异常",
)
ALT_BLOCK3_LEAF_ROW: tuple[str | None, ...] = (
    "日期", "凭证编号", "业务内容", "对方科目", "金额",
    "日期/编号", "被投资单位名称", "金额",
    "日期/编号", "付款方", "金额",
    None, None, None,
)

# G0-7 14 个叶子列
RELIABILITY_LEAF_COLUMNS: tuple[tuple[str, str], ...] = (
    ("A", "序号"),
    ("B", "函证索引号"),
    ("C", "被询证单位名称"),
    ("D", "回函方式"),
    ("E", "是否由审计项目组直接接收"),
    ("F", "是否寄回原件"),
    ("G", "被函证者身份确认（注1）"),
    ("H", "发函及回函传真信息及验证"),
    ("I", "发函邮箱"),
    ("J", "回函邮箱"),
    ("K", "邮箱可靠性验证（注2）"),
    ("L", "是否致电被函证者确认"),
    ("M", "对函证信息可靠性的考虑（注3）"),
    ("N", "回函可靠性结论"),
)
RELIABILITY_GROUP_G5 = "期末未收回原件函证可靠性验证"

# 证券差异表 17 列 / 非证券 15 列
DIFF_SECURITIES_LEAF: tuple[tuple[str, str], ...] = (
    ("A", "询证函索引号"),
    ("B", "证券名称"), ("C", "资金账号"), ("D", "开户名称"),
    ("E", "数量"), ("F", "市价（单价）"), ("G", "账面余额"),
    ("H", "数量"), ("I", "市价（单价）"), ("J", "公允价值"),
    ("K", "数量"), ("L", "市价（单价）"), ("M", "公允价值"),
    ("N", "差异原因"), ("O", "相关支持性证据"), ("P", "是否需要调账"),
    ("Q", "备注"),
)
DIFF_NONSECURITIES_LEAF: tuple[tuple[str, str], ...] = (
    ("A", "询证函索引号"),
    ("B", "被投资单位名称"), ("C", "持股比例"), ("D", "投资金额"),
    ("E", "其他投资限制/投资条款……"),
    ("F", "持股比例"), ("G", "投资金额"), ("H", "其他投资限制/投资条款……"),
    ("I", "比例"), ("J", "金额"), ("K", "其他投资限制/投资条款……"),
    ("L", "差异原因"), ("M", "相关支持性证据"), ("N", "是否需要调账"),
    ("O", "备注"),
)


# ─── fixtures ────────────────────────────────────────────────────────────────


@pytest.fixture(scope="module")
def g0_wb():
    assert _G0_PATH.exists(), f"源模板缺失: {_G0_PATH}"
    return openpyxl.load_workbook(_G0_PATH, data_only=False)


@pytest.fixture(scope="module")
def d0_wb():
    assert _D0_PATH.exists(), f"源模板缺失: {_D0_PATH}"
    return openpyxl.load_workbook(_D0_PATH, data_only=False)


def _v(ws, coord: str):
    return ws[coord].value


def _row(ws, row: int, col_from: int, count: int) -> tuple:
    return tuple(ws.cell(row, c).value for c in range(col_from, col_from + count))


# ─── Requirement 1.1 —— 10 sheet 全 visible + 名称逐字（Property 1） ──────────


class TestSheetInventory:
    def test_exactly_ten_sheets_in_order(self, g0_wb):
        assert tuple(g0_wb.sheetnames) == SHEET_NAMES

    def test_all_sheets_visible(self, g0_wb):
        """G0 无隐藏 sheet —— 与 E0（10/20 隐藏）形成对照，故 G0 无 hidden⇒skip 议题。"""
        hidden = {s: g0_wb[s].sheet_state for s in g0_wb.sheetnames if g0_wb[s].sheet_state != "visible"}
        assert hidden == {}, f"出现隐藏 sheet: {hidden}"

    def test_directory_entries_are_index_source_of_truth(self, g0_wb):
        ws = g0_wb["底稿目录"]
        got = tuple(
            (ws.cell(r, 5).value, ws.cell(r, 6).value) for r in range(4, 13)
        )
        assert got == DIRECTORY_ENTRIES

    @pytest.mark.parametrize(
        "tab_name,directory_index",
        [
            (SHEET_DIFF_SECURITIES, "G0-4"),
            (SHEET_DIFF_NONSECURITIES, "G0-5"),
            (SHEET_FRAUD, "G0-8"),
        ],
    )
    def test_tab_index_typos_are_source_facts(self, tab_name, directory_index, g0_wb):
        """三处 tab 名索引号与底稿目录不一致 —— 源模板事实，不得"修正"。"""
        assert tab_name in g0_wb.sheetnames
        assert directory_index not in tab_name, (
            f"{tab_name} 若已含 {directory_index} 说明源模板被改过，"
            "本 spec 的「定位/展示分离」设计前提失效，须复核 design.md §6"
        )

    def test_selfcheck_sheet_name_assertion_is_live(self, g0_wb):
        """反向自检：篡改任一 sheet 名常量必打红。"""
        tampered = ("底稿目录XX",) + SHEET_NAMES[1:]
        assert tuple(g0_wb.sheetnames) != tampered


# ─── Requirement 1.2 —— G0-1 上区（Property 2） ──────────────────────────────


class TestSummaryUpperZone:
    def test_r5_group_headers(self, g0_wb):
        ws = g0_wb[SHEET_SUMMARY]
        for coord, text in SUMMARY_R5_GROUPS.items():
            assert _v(ws, coord) == text, f"{coord} 期望 {text!r} 实际 {_v(ws, coord)!r}"

    def test_28_leaf_columns(self, g0_wb):
        ws = g0_wb[SHEET_SUMMARY]
        assert len(SUMMARY_LEAF_COLUMNS) == 28
        # A/B/AB 是 R5 起纵向合并（A5:A7 / B5:B7 / AB5:AB7），其余叶子在 R6
        for col, text in SUMMARY_LEAF_COLUMNS:
            row = 5 if col in ("A", "B", "AB") else 6
            assert _v(ws, f"{col}{row}") == text, f"{col}{row} 期望 {text!r} 实际 {_v(ws, f'{col}{row}')!r}"

    def test_source_template_has_no_contact_or_currency_columns(self, g0_wb):
        """源模板无 联系人/联系电话/币种 → 平台列集须剔除（Requirement 2.2）。"""
        labels = {text for _, text in SUMMARY_LEAF_COLUMNS}
        for absent in ("联系人", "联系电话", "币种"):
            assert absent not in labels

    def test_r5_merge_ranges(self, g0_wb):
        ws = g0_wb[SHEET_SUMMARY]
        got = tuple(
            sorted(str(m) for m in ws.merged_cells.ranges if m.min_row == 5 and m.max_row == 5)
        )
        assert got == tuple(sorted(SUMMARY_R5_MERGES))

    def test_audit_conclusion_is_fifth_group(self, g0_wb):
        """AB 列「审计结论」是独立段（AB5:AB7 纵向合并）→ 平台须补此列（Requirement 2.1）。"""
        ws = g0_wb[SHEET_SUMMARY]
        merged = {str(r) for r in ws.merged_cells.ranges}
        assert "AB5:AB7" in merged
        assert _v(ws, "AB5") == "审计结论"

    def test_difference_formula_handles_unreplied(self, g0_wb):
        ws = g0_wb[SHEET_SUMMARY]
        assert _v(ws, "T8") == '=IF(L8="是",S8-F8,"未回函")'

    @pytest.mark.parametrize("col,idx", sorted(SUMMARY_VLOOKUP_COL_INDEX.items()))
    def test_vlookup_column_index_from_g0_2(self, col, idx, g0_wb):
        ws = g0_wb[SHEET_SUMMARY]
        formula = _v(ws, f"{col}8")
        assert isinstance(formula, str) and formula.startswith("=VLOOKUP(")
        assert f"'{SHEET_ENTITY_VERIFY}'!A:AL,{idx},0)" in formula, formula

    def test_total_row_18(self, g0_wb):
        ws = g0_wb[SHEET_SUMMARY]
        assert _v(ws, "A18") == "合计"
        for col in ("F", "S", "T", "U", "Y", "Z"):
            assert _v(ws, f"{col}18") == f"=SUM({col}8:{col}17)"


# ─── Requirement 1.3 —— G0-1 下区四块（Property 8） ─────────────────────────


class TestSummaryLowerZone:
    @pytest.mark.parametrize("coord,text", SUMMARY_LOWER_BLOCKS)
    def test_block_anchors(self, coord, text, g0_wb):
        assert _v(g0_wb[SHEET_SUMMARY], coord) == text

    @pytest.mark.parametrize("coord,text", SUMMARY_MATRIX_METRICS)
    def test_matrix_metrics(self, coord, text, g0_wb):
        assert _v(g0_wb[SHEET_SUMMARY], coord) == text

    def test_matrix_has_eight_metrics(self):
        assert len(SUMMARY_MATRIX_METRICS) == 8

    @pytest.mark.parametrize("coord,text", SUMMARY_MATRIX_CATEGORIES)
    def test_matrix_categories(self, coord, text, g0_wb):
        assert _v(g0_wb[SHEET_SUMMARY], coord) == text

    def test_matrix_h20_is_extensible_placeholder(self):
        """H20 是 `……` 可扩位 → 品种清单可超出源模板 3 个（裁决门 A 的依据）。"""
        assert SUMMARY_MATRIX_CATEGORIES[-1] == ("H20", "……")

    def test_program_row1_enumerates_seven_varieties(self, g0_wb):
        """G0A 程序 1 明列的投资品种 —— 矩阵扩展品种的 source_ref。"""
        text = _v(g0_wb[SHEET_PROGRAM], "B7")
        for variety in (
            "交易性金融资产", "债权投资", "长期应收款", "其他债权投资",
            "长期股权投资", "其他权益工具投资", "交易性金融负债",
        ):
            assert variety in text, f"程序 1 未提及 {variety}"

    def test_matrix_formulas_are_sumif_over_category(self, g0_wb):
        """R22/R24/R27 = SUMIF(品种列 E, 品种, 金额列 F/U/Y)；R28 = (替代+回函)/账面。"""
        ws = g0_wb[SHEET_SUMMARY]
        assert _v(ws, "E22") == "=SUMIF($E$8:$E$17,E20,$F$8:$F$17)"
        assert _v(ws, "F24") == "=SUMIF($E$8:$E$17,F$20,U8:U17)"
        assert _v(ws, "E27") == "=SUMIF($E$8:$E$17,E$20,$Y$8:$Y$17)"
        assert _v(ws, "E28") == "=(E27+E24)/E21"

    def test_matrix_ratio_formulas_use_iserror_fallback(self, g0_wb):
        ws = g0_wb[SHEET_SUMMARY]
        assert _v(ws, "E23") == "=IF(ISERROR(E22/E21),0,E22/E21)"
        assert _v(ws, "E25") == "=IF(ISERROR(E24/E22),0,E24/E22)"
        assert _v(ws, "E26") == "=IF(ISERROR(E24/E21),0,E24/E21)"

    def test_book_amount_row_has_no_formula(self, g0_wb):
        """R21「本期（期末）账面金额」在源模板是手填 → 平台允许手工且手工优先。"""
        ws = g0_wb[SHEET_SUMMARY]
        for col in ("E", "F", "G"):
            v = _v(ws, f"{col}21")
            assert v is None or not (isinstance(v, str) and v.startswith("="))

    @pytest.mark.parametrize("coord,text", SUMMARY_SAMPLE_LABELS)
    def test_sample_selection_six_items(self, coord, text, g0_wb):
        assert _v(g0_wb[SHEET_SUMMARY], coord) == text

    def test_sample_selection_has_six_items(self):
        assert len(SUMMARY_SAMPLE_LABELS) == 6

    @pytest.mark.parametrize("coord,text", SUMMARY_AUDIT_NOTE_TITLES)
    def test_audit_note_five_items(self, coord, text, g0_wb):
        assert _v(g0_wb[SHEET_SUMMARY], coord) == text

    def test_split_sentence_cells_must_be_merged_when_rendered(self, g0_wb):
        """X21+X22 是一句话被拆成两格 → 平台须合并渲染，否则界面出现半句话。"""
        ws = g0_wb[SHEET_SUMMARY]
        head = _v(ws, "X21")
        tail = _v(ws, "X22")
        assert head.endswith("人民币"), head
        assert tail.startswith("（）万元"), tail

    def test_reference_conclusions_abc(self, g0_wb):
        ws = g0_wb[SHEET_SUMMARY]
        assert (_v(ws, "A55"), _v(ws, "B55")) == ("A、", "未见异常。")
        assert _v(ws, "A56") == "B、"
        assert _v(ws, "A57") == "C、"

    def test_preparation_notes_cite_cas1312_article10(self, g0_wb):
        ws = g0_wb[SHEET_SUMMARY]
        assert _v(ws, "A33") == "编制说明："
        assert "1312" in _v(ws, "A35")
        assert _v(ws, "A37") == "（一）金额较大的项目；"


# ─── Requirement 1.4 —— G0-2 38 列 3 段 ──────────────────────────────────────


class TestEntityVerify:
    def test_r5_group_headers(self, g0_wb):
        ws = g0_wb[SHEET_ENTITY_VERIFY]
        for coord, text in ENTITY_VERIFY_R5_GROUPS.items():
            assert _v(ws, coord) == text

    def test_second_send_block(self, g0_wb):
        ws = g0_wb[SHEET_ENTITY_VERIFY]
        assert _v(ws, "AC5") == "经核实的原因"
        assert _v(ws, "AD5") == "该原因是否合理（是/否）\n（注3）"
        assert _v(ws, "AE5") == "是否进行第二次发函？\n（是/否）"
        assert _v(ws, "AF5") == "第二次发函的被函证单位信息"
        assert _v(ws, "AL5") == "第二次发函结果？（送抵/退回）\n（注4）"

    @pytest.mark.parametrize("col,text", ENTITY_VERIFY_LEAF_R6)
    def test_leaf_columns_r6(self, col, text, g0_wb):
        assert _v(g0_wb[SHEET_ENTITY_VERIFY], f"{col}6") == text

    def test_last_header_is_al_and_beyond_is_empty(self, g0_wb):
        """表格区 A..AL = 38 列；`max_column` 为 43 是尾部残留格式，非数据列。"""
        ws = g0_wb[SHEET_ENTITY_VERIFY]
        assert _v(ws, "AL5") == "第二次发函结果？（送抵/退回）\n（注4）"
        for col in ("AM", "AN", "AO", "AP", "AQ", "AR", "AS"):
            assert _v(ws, f"{col}5") is None
            assert _v(ws, f"{col}6") is None


# ─── Requirement 1.5 —— G0-3 / F0-8 与 D0 逐字相同（Property 25） ────────────


class TestSharedWithD0:
    def test_followup_memo_identical_to_d0(self, g0_wb, d0_wb):
        g = g0_wb[SHEET_FOLLOWUP]
        d = d0_wb["跟函函证过程控制D0-3"]
        # A5..A31 是备忘录正文 + 3 核对点 + 签名 + 提示（A3/A4 是引用底稿目录的公式，跳过）
        g_text = [g.cell(r, 1).value for r in range(5, 32)]
        d_text = [d.cell(r, 1).value for r in range(5, 32)]
        assert g_text == d_text

    def test_fraud_items_identical_to_d0(self, g0_wb, d0_wb):
        g = g0_wb[SHEET_FRAUD]
        d = d0_wb["函证程序舞弊风险评价表D0-8"]
        g_items = [g.cell(r, 1).value for r in range(6, 26)]
        d_items = [d.cell(r, 1).value for r in range(6, 26)]
        assert g_items == d_items
        assert len([x for x in g_items if x and x != "……"]) == 19

    def test_fraud_pushes_to_b50(self, g0_wb):
        assert _v(g0_wb[SHEET_FRAUD], "H26") == "B50"

    def test_followup_three_check_points(self, g0_wb):
        ws = g0_wb[SHEET_FOLLOWUP]
        assert _v(ws, "A23") == "是否了解处理函证的通常流程和处理人员"
        assert _v(ws, "A24") == "是否确认询证函处理人员的身份及权限"
        assert _v(ws, "A25") == "处理人员是否按正常流程处理"


# ─── Requirement 1.6 —— G0A 12 条 + 程序分类（Property 19） ──────────────────


class TestProgramTable:
    def test_twelve_programs(self):
        assert len(PROGRAM_ROWS) == 12

    @pytest.mark.parametrize("seq,category,ref_index", PROGRAM_ROWS)
    def test_program_row(self, seq, category, ref_index, g0_wb):
        ws = g0_wb[SHEET_PROGRAM]
        row = 6 + int(seq)
        assert ws.cell(row, 1).value == seq, "A 列 seq 是字符串不是整数"
        assert ws.cell(row, 4).value == category
        assert ws.cell(row, 5).value == ref_index

    def test_category_column_header(self, g0_wb):
        ws = g0_wb[SHEET_PROGRAM]
        assert _v(ws, "D5") == "程序分类"
        assert _v(ws, "E5") == "底稿索引号"

    def test_optional_programs_exist(self):
        """备选（7、8）与 IPO 专项（2、11）—— 平台 applicable_default 不得为 yes。"""
        optional = {seq for seq, cat, _ in PROGRAM_ROWS if cat == "备选"}
        ipo = {seq for seq, cat, _ in PROGRAM_ROWS if "IPO" in cat}
        assert optional == {"7", "8"}
        assert ipo == {"2", "11"}

    def test_all_categories_non_empty(self):
        assert all(cat for _, cat, _ in PROGRAM_ROWS)

    def test_two_programs_have_no_ref_index(self):
        assert {seq for seq, _, ref in PROGRAM_ROWS if ref is None} == {"8", "12"}


# ─── Requirement 1.7 —— G0-6 区块（Property 16） ─────────────────────────────


class TestAlternativeProcedure:
    @pytest.mark.parametrize("coord,text", ALT_HEADER_FIELDS)
    def test_header_fields(self, coord, text, g0_wb):
        assert _v(g0_wb[SHEET_ALTERNATIVE], coord) == text

    @pytest.mark.parametrize("coord,text", ALT_BLOCK_ANCHORS)
    def test_block_anchors(self, coord, text, g0_wb):
        assert _v(g0_wb[SHEET_ALTERNATIVE], coord) == text

    @pytest.mark.parametrize("coord,text", ALT_SCOPE_CELLS)
    def test_scope_option_cells(self, coord, text, g0_wb):
        """两组 5 点选项逐字。`B7` 末项「全部」/ `C15` 末项「其他」，只差最后一项。"""
        assert _v(g0_wb[SHEET_ALTERNATIVE], coord) == text

    def test_two_scope_option_sets_differ_only_in_last_item(self):
        """反向自检：若有人把两组选项抄成一样，本断言打红。"""
        a = ALT_TEST_SCOPE_TEXT.replace("（）", "|").strip("|").split("|")
        b = ALT_OCCURRENCE_SCOPE_TEXT.replace("（）", "|").strip("|").split("|")
        assert len(a) == len(b) == 5
        assert a[:4] == b[:4]
        assert (a[4], b[4]) == ("全部", "其他")

    @pytest.mark.parametrize("coord,text", ALT_NUMBERED_SECTIONS)
    def test_numbered_sections(self, coord, text, g0_wb):
        """源模板恰四个带序号的段，且「二」是检查过程记录。

        平台改造前用源外增强段「余额汇总与检查比例」占了「二」，导致
        检查过程记录被编成「三」、审计说明与结论被挤成合并的「四」。
        前端交叉锁死见 `alternativeG06/__tests__/g06SourceFidelity.spec.ts`。
        """
        assert _v(g0_wb[SHEET_ALTERNATIVE], coord) == text

    def test_no_balance_summary_section_in_source(self, g0_wb):
        """反向自检：源模板 G0-6 里**不存在**「余额汇总」类段（它是平台源外增强）。"""
        ws = g0_wb[SHEET_ALTERNATIVE]
        texts = [
            str(ws.cell(row=r, column=c).value or "")
            for r in range(1, ws.max_row + 1)
            for c in range(1, ws.max_column + 1)
        ]
        joined = "".join(texts)
        assert "余额汇总" not in joined
        # 而「检查过程记录」确实在源模板里 → 证明上面这条不是因为读不到内容而空过
        assert "检查过程记录" in joined

    @pytest.mark.parametrize("coord,text", ALT_PREPARATION_NOTES)
    def test_preparation_notes(self, coord, text, g0_wb):
        """三条替代程序要点在 **B 列**（`A46`/`A47` 只是标题）—— 只扫 A 列会整段漏掉。"""
        assert _v(g0_wb[SHEET_ALTERNATIVE], coord) == text

    def test_block1_has_no_voucher_columns(self, g0_wb):
        """区块① 源模板只有 5 列且**无记账凭证列** → 平台现有实现塞了 5 个凭证列，须删。"""
        ws = g0_wb[SHEET_ALTERNATIVE]
        assert _row(ws, 10, 1, 5) == ALT_BLOCK1_COLUMNS
        assert "记账凭证" not in (ALT_BLOCK1_COLUMNS)
        assert "投资条款" in ALT_BLOCK1_COLUMNS

    @pytest.mark.parametrize("group_row,leaf_row", [(17, 18), (25, 26)])
    def test_block2_two_level_header(self, group_row, leaf_row, g0_wb):
        ws = g0_wb[SHEET_ALTERNATIVE]
        assert _row(ws, group_row, 1, 14) == ALT_BLOCK2_GROUP_ROW
        assert _row(ws, leaf_row, 1, 14) == ALT_BLOCK2_LEAF_ROW

    def test_block3_two_level_header(self, g0_wb):
        ws = g0_wb[SHEET_ALTERNATIVE]
        assert _row(ws, 33, 1, 14) == ALT_BLOCK3_GROUP_ROW
        assert _row(ws, 34, 1, 14) == ALT_BLOCK3_LEAF_ROW

    def test_block2_support_docs_are_six_columns(self):
        """支持性文件 1/2 各 3 列 —— 平台把它压成一个 `support_doc` 单列，须拆。"""
        leaves = [x for x in ALT_BLOCK2_LEAF_ROW[5:11]]
        assert leaves == ["识别特征", "信息1", "信息2", "识别特征", "信息1", "信息2"]


# ─── Requirement 1.5 补充 —— G0-7 14 列（Property 26 的源侧） ────────────────


class TestReliability:
    def test_group_header(self, g0_wb):
        assert _v(g0_wb[SHEET_RELIABILITY], "G5") == RELIABILITY_GROUP_G5

    @pytest.mark.parametrize("col,text", RELIABILITY_LEAF_COLUMNS)
    def test_leaf_columns(self, col, text, g0_wb):
        ws = g0_wb[SHEET_RELIABILITY]
        # A..F 与 N 是 R5 起纵向合并（A5:A6 等），G..M 在 R6
        row = 6 if col in ("G", "H", "I", "J", "K", "L", "M") else 5
        assert _v(ws, f"{col}{row}") == text

    def test_fourteen_leaf_columns(self):
        assert len(RELIABILITY_LEAF_COLUMNS) == 14


# ─── 两张差异核对表（归档 spec 已锁列集，此处只固化源事实） ──────────────────


class TestDiffTables:
    def test_securities_17_columns(self, g0_wb):
        ws = g0_wb[SHEET_DIFF_SECURITIES]
        assert len(DIFF_SECURITIES_LEAF) == 17
        for col, text in DIFF_SECURITIES_LEAF:
            row = 5 if col in ("A", "Q") else 6
            assert _v(ws, f"{col}{row}") == text

    def test_securities_three_dim_groups(self, g0_wb):
        ws = g0_wb[SHEET_DIFF_SECURITIES]
        assert _v(ws, "B5") == "账面结存证券投资①"
        assert _v(ws, "H5") == "证券投资回函②"
        assert _v(ws, "K5") == "差异③=①-②"

    def test_nonsecurities_15_columns(self, g0_wb):
        ws = g0_wb[SHEET_DIFF_NONSECURITIES]
        assert len(DIFF_NONSECURITIES_LEAF) == 15
        for col, text in DIFF_NONSECURITIES_LEAF:
            row = 5 if col in ("A", "O") else 6
            assert _v(ws, f"{col}{row}") == text

    def test_nonsecurities_three_dim_groups(self, g0_wb):
        ws = g0_wb[SHEET_DIFF_NONSECURITIES]
        assert _v(ws, "B5") == "账面非证券投资①"
        assert _v(ws, "F5") == "非证券投资回函②"
        assert _v(ws, "I5") == "差异③=①-②"


# ─── Requirement 10 —— 源模板自身缺陷（Property 24 的 (a) 半 / Property 27） ──


class TestSourceDefectsExist:
    """这些断言证明"缺陷确实存在"。平台侧"已按意图纠正"的断言在前端守卫里
    （`g0SourceDefects.spec.ts`），两侧缺一即视为未完成（Requirement 10.4）。"""

    def test_defect_directory_serial_hardcoded(self, g0_wb):
        """`底稿目录!D7` 硬写 2（其余为 `=D(n-1)+1`）→ 序号显示 1,2,3,2,3,4,3,4,5。"""
        ws = g0_wb["底稿目录"]
        assert _v(ws, "D4") == 1
        assert _v(ws, "D5") == "=D4+1"
        assert _v(ws, "D6") == "=D5+1"
        assert _v(ws, "D7") == 2, "D7 应为硬写的 2（缺陷）而非公式"
        assert _v(ws, "D10") == 3, "D10 亦为硬写的 3"
        rendered = [1, 2, 3, 2, 3, 4, 3, 4, 5]
        assert rendered[3] < rendered[2], "序号列确实回退（缺陷可见）"

    def test_defect_matrix_sumif_row_overflow(self, g0_wb):
        """`G0-1!E24` 的 SUMIF 求和区 `U8:U179` 越界（数据区只到 R17）。"""
        formula = _v(g0_wb[SHEET_SUMMARY], "E24")
        assert "U8:U179" in formula, formula
        # 同一行的 F24 是正确的 U8:U17 → 证明这是笔误而非有意
        assert "U8:U17)" in _v(g0_wb[SHEET_SUMMARY], "F24")

    def test_defect_xref_reliability_points_to_g0_6(self, g0_wb):
        """`G0-1!S24` 写「（G0-6）」，而回函可靠性验证是 G0-7。"""
        text = _v(g0_wb[SHEET_SUMMARY], "S24")
        assert "（G0-6）" in text
        assert "可靠性" in text

    def test_defect_entity_verify_self_reference(self, g0_wb):
        """`G0-2!AA6` 写「跟函函证控制过程（G0-2）」自引用，应为 G0-3。"""
        text = _v(g0_wb[SHEET_ENTITY_VERIFY], "AA6")
        assert text == "跟函函证控制过程（G0-2）"
        assert "G0-3" not in text

    def test_defect_securities_market_value_diff_direction(self, g0_wb):
        """证券差异表 M 列 `=J−G`（回函−账面）与表头 `③=①−②` 及 K/L 方向相反。"""
        ws = g0_wb[SHEET_DIFF_SECURITIES]
        assert _v(ws, "K7") == "=E7-H7", "数量差异 = 账面 − 回函"
        assert _v(ws, "L7") == "=F7-I7", "单价差异 = 账面 − 回函"
        assert _v(ws, "M7") == "=J7-G7", "公允价值差异写成了 回函 − 账面（缺陷）"

    def test_defect_reply_amount_group_header_shifted(self, g0_wb, d0_wb):
        """缺陷 #7：G0-1 的「3、回函金额确认」段头落在 `U5:W5`，
        而 D0-1 / F0-1 同构表均为 `S5:W5` → G0-1 的 S 回函金额 / T 差异 两列无段归属。

        平台按 D0/F0 的正确意图把 S/T/U 归入「回函金额确认」段（Requirement 10.2）。
        """
        g = g0_wb[SHEET_SUMMARY]
        assert _v(g, "S5") is None, "S5 无段头（缺陷）"
        assert _v(g, "T5") is None
        assert _v(g, "U5") == "3、回函金额确认"
        g_merges = tuple(
            sorted(str(m) for m in g.merged_cells.ranges if m.min_row == 5 and m.max_row == 5)
        )
        assert g_merges == tuple(sorted(SUMMARY_R5_MERGES))

        # 意图旁证：D0-1 同位段头在 S5:W5
        d = d0_wb["函证结果汇总表D0-1"]
        assert _v(d, "S5") == "3、回函金额确认"
        d_merges = tuple(
            sorted(str(m) for m in d.merged_cells.ranges if m.min_row == 5 and m.max_row == 5)
        )
        assert d_merges == tuple(sorted(SUMMARY_R5_MERGES_D0F0))
        assert g_merges != d_merges, "两表段头合并区应当不同（否则缺陷已不存在，须复核 spec）"

    def test_defect_alt_total_sums_index_column(self, g0_wb):
        """G0-6 借方块合计行对**索引号**列求和；贷方块无对应 M31 → 复制残留。"""
        ws = g0_wb[SHEET_ALTERNATIVE]
        assert _v(ws, "A23") == "合计"
        assert _v(ws, "E23") == "=SUM(E19:E22)"
        assert _v(ws, "M23") == "=SUM(M19:M22)", "借方块确实对索引号列求和（缺陷）"
        # 贷方块是正确形态：只有金额列合计
        assert _v(ws, "A31") == "合计"
        assert _v(ws, "E31") == "=SUM(E27:E30)"
        assert _v(ws, "M31") is None, "贷方块无 M31 → 证明 M23 是笔误而非有意"
        # M 列确实是索引号（不是金额）
        assert _v(ws, "M17") == "索引号"

    def test_defect_alt_abnormal_dv_offset(self, g0_wb):
        """「是否异常」在 N 列，但 `√,×` 验证挂在 O35:O38（右移一列且只覆盖区块③）。"""
        ws = g0_wb[SHEET_ALTERNATIVE]
        assert _v(ws, "N33") == "是否异常"
        # O 列在两级表头里无标题 → 验证挂在无标题列上
        assert _v(ws, "O33") is None
        assert _v(ws, "O34") is None

        dvs = list(ws.data_validations.dataValidation)
        # 逐格判定（memory 铁律：禁按打印顺序取第一个 sqref）
        yes_no = [dv for dv in dvs if dv.formula1 and "√" in str(dv.formula1)]
        assert len(yes_no) == 1, f"应恰有一处 √/× 验证，实为 {len(yes_no)}"
        dv = yes_no[0]
        assert "O35" in str(dv.sqref)
        # 区块② 的同名列完全没有验证 → 源模板自身未统一
        assert all("N17" not in str(d.sqref) and "N25" not in str(d.sqref) for d in dvs)

    def test_defect_tab_index_typos_count(self, g0_wb):
        """三处 tab 名索引号与底稿目录不一致。"""
        dir_index = {name: idx for name, idx in DIRECTORY_ENTRIES}
        mismatches: list[str] = []
        # 显式列出三条（按正则从 tab 名抽索引号易脆：G0A 无连字符、两张差异表索引号在名字中间）
        for tab, expected in (
            (SHEET_DIFF_SECURITIES, "G0-4"),
            (SHEET_DIFF_NONSECURITIES, "G0-5"),
            (SHEET_FRAUD, "G0-8"),
        ):
            assert expected not in tab
            mismatches.append(tab)
        assert len(mismatches) == 3
        # 底稿目录侧确实是修正后的索引号
        assert dir_index["函证差异核对表（证券投资）"] == "G0-4"
        assert dir_index["函证差异核对表(非证券投资)"] == "G0-5"
        assert dir_index["函证程序舞弊风险评价表"] == "G0-8"


# ─── 反向自检（Requirement 1.8） ─────────────────────────────────────────────


class TestReverseSelfChecks:
    def test_selfcheck_leaf_column_count_is_asserted(self):
        """若有人把 28 列常量改短，`test_28_leaf_columns` 的长度断言必打红。"""
        assert len(SUMMARY_LEAF_COLUMNS) == 28
        assert len({col for col, _ in SUMMARY_LEAF_COLUMNS}) == 28

    def test_selfcheck_wrong_metric_text_fails(self, g0_wb):
        ws = g0_wb[SHEET_SUMMARY]
        assert _v(ws, "C21") != "本期账面金额"  # 少了「（期末）」与冒号

    def test_selfcheck_matrix_formula_not_plain_sum(self, g0_wb):
        """矩阵是按品种 SUMIF 而非整列 SUM —— 去掉品种过滤即为错。"""
        ws = g0_wb[SHEET_SUMMARY]
        assert _v(ws, "E22") != "=SUM($F$8:$F$17)"

    def test_selfcheck_d0_comparison_is_live(self, g0_wb, d0_wb):
        """若两个工作簿被拿成同一个，`test_*_identical_to_d0` 会退化为自证。"""
        assert g0_wb is not d0_wb
        assert _v(g0_wb[SHEET_SUMMARY], "A2") == "投资循环函证结果汇总表"
        assert _v(d0_wb["函证结果汇总表D0-1"], "A2") != "投资循环函证结果汇总表"

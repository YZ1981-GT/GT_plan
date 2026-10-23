"""上年 Word 导入 / 上年继承（spec disclosure-note-template-word-realignment Task 15，design §5）。

路由前缀带 ``{project_id}``，原来五个端点一个鉴权依赖都没挂（未登录可读写任意项目附注，鉴权基线扫描器只看
装饰器路径所以没看见）；导入服务读不存在的 ``note.title``、调不存在的 ``_extract_to_opening_balance``、年度恒为
2025。真 SQLite + 真 ORM + 真路由与真项目鉴权。
"""
from __future__ import annotations

import io
import json
import uuid
from collections.abc import AsyncGenerator
from datetime import datetime, timedelta, timezone

import pytest
import pytest_asyncio
from fastapi import FastAPI
from httpx import ASGITransport, AsyncClient
from sqlalchemy.ext.asyncio import AsyncSession, async_sessionmaker, create_async_engine

import tests.conftest  # noqa: F401  注册全部模型 + SQLite JSONB 兼容
from app.core.database import get_db
from app.deps import get_current_user
from app.models.base import Base, PermissionLevel, ProjectStatus, ProjectUserRole
from app.models.core import Project, ProjectUser, User
from app.models.report_models import DisclosureNote
from app.models.staff_models import ProjectAssignment, StaffMember
from app.routers import note_advanced
from tests.test_note_text_batch_endpoints import _Actor

YEAR = 2025
OLD = datetime(2020, 1, 1, tzinfo=timezone.utc)
HEADERS = ["项目", "期末余额", "期初余额"]


def _table(values: list[list]) -> dict:
    rows = [{"label": label, "values": list(v)} for label, v in zip(("库存现金", "银行存款"), values)]
    return {"name": "货币资金", "headers": HEADERS, "rows": rows}


def _docx(sections: list[tuple[str, str, list[list[str]] | None]]) -> bytes:
    import docx

    doc = docx.Document()
    for title, para, rows in sections:
        doc.add_heading(title, level=2)
        if para:
            doc.add_paragraph(para)
        if rows:
            t = doc.add_table(rows=len(rows), cols=len(rows[0]))
            for i, row in enumerate(rows):
                for j, value in enumerate(row):
                    t.cell(i, j).text = value
    buf = io.BytesIO()
    doc.save(buf)
    return buf.getvalue()


PRIOR_DOCX = _docx([
    ("货币资金", "上年货币资金说明正文。", [HEADERS, ["库存现金", "1,234.50", "800"], ["银行存款", "(66)", "9"]]),
    ("会计期间", "上年会计期间说明。", None),
])


@pytest_asyncio.fixture
async def env() -> AsyncGenerator[dict, None]:
    engine = create_async_engine("sqlite+aiosqlite:///:memory:")
    async with engine.begin() as conn:
        await conn.run_sync(lambda c: Base.metadata.create_all(
            c, tables=[User.__table__, Project.__table__, ProjectUser.__table__, DisclosureNote.__table__,
                         StaffMember.__table__, ProjectAssignment.__table__]))
    factory = async_sessionmaker(engine, class_=AsyncSession, expire_on_commit=False)
    ids = {k: uuid.uuid4() for k in ("me", "viewer", "p", "cash", "period", "prior_cash", "prior_gone")}
    first = _table([[None, None], [5, 7]])
    async with factory() as s:
        s.add(Project(id=ids["p"], name="甲_2025", client_name="甲公司", template_type="soe",
                      report_scope="standalone", status=ProjectStatus.execution))
        await s.flush()
        s.add_all([
            ProjectUser(project_id=ids["p"], user_id=ids["me"], role=ProjectUserRole.auditor,
                        permission_level=PermissionLevel.edit),
            ProjectUser(project_id=ids["p"], user_id=ids["viewer"], role=ProjectUserRole.auditor),
            DisclosureNote(id=ids["cash"], project_id=ids["p"], year=YEAR, note_section="八、1", section_title="货币资金",
                           table_data={"_tables": [first], **_table([[None, None], [5, 7]])},
                           updated_at=OLD, updated_by=None),
            DisclosureNote(id=ids["period"], project_id=ids["p"], year=YEAR, note_section="四、会计期间",
                           section_title="会计期间", text_content="<p>内容模板写的</p>", is_local_override=True,
                           updated_at=OLD, updated_by=None),
            DisclosureNote(id=ids["prior_cash"], project_id=ids["p"], year=YEAR - 1, note_section="八、1",
                           section_title="货币资金", text_content="上年库里的货币资金说明"),
            DisclosureNote(id=ids["prior_gone"], project_id=ids["p"], year=YEAR - 1, note_section="四、会计期间",
                           section_title="会计期间", text_content="回收站里的不该被继承", is_deleted=True),
        ])
        await s.commit()

    app = FastAPI()
    app.include_router(note_advanced.router)

    async def _override_db():
        async with factory() as session:
            yield session

    app.dependency_overrides[get_db] = _override_db
    actor = {"user": _Actor(ids["me"])}
    app.dependency_overrides[get_current_user] = lambda: actor["user"]
    async with AsyncClient(transport=ASGITransport(app=app), base_url="http://test") as client:
        yield {"client": client, "ids": ids, "factory": factory, "actor": actor, "app": app}
    await engine.dispose()


async def _note(env: dict, key: str) -> DisclosureNote:
    async with env["factory"]() as db:
        return await db.get(DisclosureNote, env["ids"][key])


def _utc(value: datetime) -> datetime:
    return value.replace(tzinfo=timezone.utc) if value.tzinfo is None else value.astimezone(timezone.utc)


async def _import(env: dict, year: int | None = YEAR):
    url = f"/api/projects/{env['ids']['p']}/notes/import-prior-year" + (f"?year={year}" if year else "")
    files = {"file": ("上年附注.docx", PRIOR_DOCX,
                      "application/vnd.openxmlformats-officedocument.wordprocessingml.document")}
    return await env["client"].post(url, files=files)


@pytest.mark.asyncio
async def test_every_endpoint_requires_login(env):
    """不带身份 ⇒ 401/403（原来五个端点全部 200）。"""
    env["app"].dependency_overrides.pop(get_current_user)
    base = f"/api/projects/{env['ids']['p']}/notes"
    responses = [
        await _import(env),
        await env["client"].post(f"{base}/inherit-prior-year?year={YEAR}"),
        await env["client"].get(f"{base}/cross-references?year={YEAR}"),
        await env["client"].post(f"{base}/cross-references/update?year={YEAR}"),
        await env["client"].post(f"{base}/generate-variation-analysis?year={YEAR}"),
    ]
    assert [r.status_code in (401, 403) for r in responses] == [True] * 5, [r.status_code for r in responses]


@pytest.mark.asyncio
async def test_readonly_member_cannot_write(env):
    env["actor"]["user"] = _Actor(env["ids"]["viewer"])
    assert (await _import(env)).status_code == 403
    assert (await env["client"].post(
        f"/api/projects/{env['ids']['p']}/notes/inherit-prior-year?year={YEAR}")).status_code == 403
    assert (await _note(env, "cash")).updated_by is None


@pytest.mark.asyncio
async def test_year_is_required(env):
    assert (await _import(env, year=None)).status_code == 422


@pytest.mark.asyncio
async def test_import_fills_text_and_opening_column_and_stamps(env):
    since = datetime.now(timezone.utc) - timedelta(seconds=1)
    resp = await _import(env)
    assert resp.status_code == 200, resp.text
    assert resp.json()["matched"] == 2
    assert (resp.json()["filled_cells"], resp.json()["pending_tables"], resp.json()["no_source_tables"]) == (1, 0, 0)
    assert resp.json()["changed_sections"] == ["八、1"], "只报真被改写的章节（内容模板写过的会计期间没动）"
    cash, period = await _note(env, "cash"), await _note(env, "period")
    assert cash.text_content == "上年货币资金说明正文。"
    # 期初列（下标 1）：空格填上年期末（千分位 / 括号负数照解析）；已有值的不覆盖
    assert [r["values"] for r in cash.table_data["_tables"][0]["rows"]] == [[None, 1234.5], [5, 7]]
    assert [r["values"] for r in cash.table_data["rows"]] == [[None, 1234.5], [5, 7]]  # 顶层镜像一并填
    assert cash.template_lineage["prior_year_status"] == "pending_update"
    assert cash.updated_by == env["ids"]["me"] and _utc(cash.updated_at) >= since
    # 内容模板写过的正文不覆盖，也不算编辑
    assert (period.text_content, period.updated_by, _utc(period.updated_at)) == ("<p>内容模板写的</p>", None, OLD)


async def _post(env: dict, mode: str, *, data: dict | None = None, file: bytes = PRIOR_DOCX):
    url = f"/api/projects/{env['ids']['p']}/notes/import-prior-year?year={YEAR}&mode={mode}"
    files = {"file": ("上年附注.docx", file, "application/vnd.openxmlformats-officedocument.wordprocessingml.document")}
    return await env["client"].post(url, files=files, data=data or {})


@pytest.mark.asyncio
async def test_preview_plans_each_table_without_writing(env):
    """预览（R6.6）：逐表给出状态、上年表、列对应与可填格数；不写库。"""
    resp = await _post(env, "preview")
    assert resp.status_code == 200, resp.text
    body = resp.json()
    assert body["status"] == "preview" and body["counts"] == {"auto": 1, "needs_confirm": 0, "no_source": 0}
    [table] = body["tables"]
    assert (table["note_section"], table["status"], table["prior_table_id"], table["col_map"], table["fillable"]) == (
        "八、1", "auto", "t1", {"2": 1}, 1)
    assert table["candidate_fillable"] == {"t1": 1}, "库存现金期初是空格可填；银行存款期初已有 7，不覆盖"
    assert [c["label"] for c in table["opening_columns"]] == ["期初余额"]
    assert [(t["id"], t["section"], [c["label"] for c in t["columns"]]) for t in body["prior_tables"]] == [
        ("t1", "货币资金", ["期末余额", "期初余额"])]
    cash = await _note(env, "cash")
    assert [r["values"] for r in cash.table_data["_tables"][0]["rows"]] == [[None, None], [5, 7]]
    assert (cash.text_content, cash.updated_by) == (None, None)


@pytest.mark.asyncio
async def test_preview_gives_fillable_only_for_candidates_with_columns(env):
    """候选表一列都对不上 ⇒ 不给可填格数（原来给 0，弹窗显示「可填 0 格」像是行名对不上）。"""
    two = _docx([("货币资金", "", [HEADERS, ["库存现金", "1", "2"]])])
    import docx as _d

    doc = _d.Document(io.BytesIO(two))
    t = doc.add_table(rows=2, cols=3)
    for i, row in enumerate([["项目", "本期增加", "本期减少"], ["库存现金", "5", "6"]]):
        for j, value in enumerate(row):
            t.cell(i, j).text = value
    buf = io.BytesIO()
    doc.save(buf)
    [table] = (await _post(env, "preview", file=buf.getvalue())).json()["tables"]
    assert table["candidates"] == ["t1", "t2"] and table["candidate_maps"]["t2"] == {}
    assert table["candidate_fillable"] == {"t1": 1}


@pytest.mark.asyncio
async def test_apply_writes_only_what_was_confirmed(env):
    preview = (await _post(env, "preview")).json()
    [table] = preview["tables"]
    sha = preview["file_sha256"]
    skip = [{"key": table["key"], "prior_table_id": None}]
    resp = await _post(env, "apply", data={"decisions": json.dumps(skip), "file_sha256": sha, "import_text": "false"})
    assert resp.status_code == 200 and (resp.json()["filled_cells"], resp.json()["changed_sections"]) == (0, [])
    assert (await _note(env, "cash")).updated_by is None, "不导入、不填正文 ⇒ 不算编辑"
    # 弹窗里改选列：期初列取上年表的「期初余额」列
    pick = [{"key": table["key"], "prior_table_id": table["prior_table_id"], "col_map": {"2": 2}}]
    resp = await _post(env, "apply", data={"decisions": json.dumps(pick), "file_sha256": sha})
    assert resp.status_code == 200, resp.text
    assert (resp.json()["filled_cells"], resp.json()["text_filled"], resp.json()["changed_sections"]) == (1, 1, ["八、1"])
    cash = await _note(env, "cash")
    assert [r["values"] for r in cash.table_data["_tables"][0]["rows"]] == [[None, 800.0], [5, 7]]
    assert [r["values"] for r in cash.table_data["rows"]] == [[None, 800.0], [5, 7]]  # 顶层镜像一并填
    assert cash.text_content == "上年货币资金说明正文。" and cash.updated_by == env["ids"]["me"]


@pytest.mark.asyncio
async def test_apply_rejects_another_file_or_bad_decisions(env):
    preview = (await _post(env, "preview")).json()
    [table] = preview["tables"]
    ok = {"file_sha256": preview["file_sha256"]}
    other = _docx([("货币资金", "改过的正文", [HEADERS, ["库存现金", "1", "2"]])])
    resp = await _post(env, "apply", data={**ok, "decisions": "[]"}, file=other)
    assert resp.status_code == 409 and "重新预览" in resp.json()["detail"]
    closing = [{"key": table["key"], "prior_table_id": "t1", "col_map": {"1": 1}}]  # 第 1 列是期末列，不收
    resp = await _post(env, "apply", data={**ok, "decisions": json.dumps(closing)})
    assert resp.status_code == 400 and "列对应不合法" in resp.json()["detail"]
    for bad in ([{"key": "x:0", "prior_table_id": "t1"}], [{"key": table["key"], "prior_table_id": "t9"}]):
        assert (await _post(env, "apply", data={**ok, "decisions": json.dumps(bad)})).status_code == 400
    assert (await _post(env, "apply", data={**ok, "decisions": "不是 JSON"})).status_code == 400
    assert (await _post(env, "apply", data={"decisions": "[]"})).status_code == 400, "没带文件指纹"
    assert (await _note(env, "cash")).updated_by is None


@pytest.mark.asyncio
async def test_extension_check_is_case_insensitive(env):
    """前端按扩展名大小写不敏感放行 ⇒ 后端同口径（原来「上年附注.DOCX」被 400「仅支持 .docx」）；别的格式仍拒。"""
    url = f"/api/projects/{env['ids']['p']}/notes/import-prior-year?year={YEAR}&mode=preview"
    for name, status in (("上年附注.DOCX", 200), ("上年附注.doc", 400)):
        resp = await env["client"].post(url, files={"file": (name, PRIOR_DOCX, "application/octet-stream")})
        assert resp.status_code == status, (name, resp.text)


@pytest.mark.asyncio
async def test_reimport_with_nothing_new_does_not_stamp(env):
    assert (await _import(env)).status_code == 200
    async with env["factory"]() as db:
        note = await db.get(DisclosureNote, env["ids"]["cash"])
        note.updated_at, note.updated_by = OLD, None
        await db.commit()
    assert (await _import(env)).status_code == 200
    cash = await _note(env, "cash")
    assert (_utc(cash.updated_at), cash.updated_by) == (OLD, None)


@pytest.mark.asyncio
async def test_inherit_copies_prior_text_stamps_and_ignores_recycle_bin(env):
    since = datetime.now(timezone.utc) - timedelta(seconds=1)
    resp = await env["client"].post(f"/api/projects/{env['ids']['p']}/notes/inherit-prior-year?year={YEAR}")
    assert resp.status_code == 200, resp.text
    assert resp.json()["inherited"] == 1
    cash, period = await _note(env, "cash"), await _note(env, "period")
    assert cash.text_content == "上年库里的货币资金说明"
    assert cash.updated_by == env["ids"]["me"] and _utc(cash.updated_at) >= since
    assert cash.template_lineage["prior_year_source"] == YEAR - 1
    assert period.text_content == "<p>内容模板写的</p>" and period.updated_by is None


# ── R7.2 合并锁定时拒绝；R7.3 上年正文只是占位话不导入（spec consol-note-header-library-and-node-reaggregate） ──


async def _lock(env: dict) -> None:
    async with env["factory"]() as db:
        project = await db.get(Project, env["ids"]["p"])
        project.consol_lock = True
        await db.commit()


@pytest.mark.asyncio
async def test_consol_lock_rejects_import_preview_and_inherit(env):
    """合并锁定 ⇒ 上年导入（含预览）与上年继承 423，库里一格不动（原来 consol_lock=False 照写）。"""
    await _lock(env)
    base = f"/api/projects/{env['ids']['p']}/notes"
    responses = [await _post(env, "preview"), await _post(env, "auto"),
                 await env["client"].post(f"{base}/inherit-prior-year?year={YEAR}")]
    assert [r.status_code for r in responses] == [423, 423, 423], [r.text for r in responses]
    assert {r.json()["detail"] for r in responses} == {"项目已被合并锁定，无法修改"}
    cash = await _note(env, "cash")
    assert (cash.text_content, cash.updated_by) == (None, None)
    assert [r["values"] for r in cash.table_data["_tables"][0]["rows"]] == [[None, None], [5, 7]]


PLACEHOLDER_DOCX = _docx([
    ("货币资金", "本期无此项业务。", [HEADERS, ["库存现金", "1,234.50", "800"], ["银行存款", "(66)", "9"]]),
    ("会计期间", "不适用。", None),
])


@pytest.mark.asyncio
async def test_placeholder_prior_text_is_not_imported(env):
    """上年正文只是「本期无此项业务。」⇒ 不覆盖本年正文（表格照填）；预览的「将导入正文」章节数不算它。"""
    async with env["factory"]() as db:
        note = await db.get(DisclosureNote, env["ids"]["cash"])
        note.text_content = "本年已写的货币资金说明。"
        await db.commit()
    preview = (await _post(env, "preview", file=PLACEHOLDER_DOCX)).json()
    assert preview["text_sections"] == 0
    resp = await _post(env, "auto", file=PLACEHOLDER_DOCX)
    assert resp.status_code == 200, resp.text
    assert (resp.json()["text_filled"], resp.json()["filled_cells"]) == (0, 1)
    cash = await _note(env, "cash")
    assert cash.text_content == "本年已写的货币资金说明。"
    assert [r["values"] for r in cash.table_data["_tables"][0]["rows"]] == [[None, 1234.5], [5, 7]]
    # 按确认写入同口径：勾了表、正文照样不覆盖
    [table] = preview["tables"]
    pick = [{"key": table["key"], "prior_table_id": table["prior_table_id"], "col_map": table["col_map"]}]
    resp = await _post(env, "apply", file=PLACEHOLDER_DOCX,
                       data={"decisions": json.dumps(pick), "file_sha256": preview["file_sha256"]})
    assert resp.status_code == 200 and resp.json()["text_filled"] == 0, resp.text
    assert (await _note(env, "cash")).text_content == "本年已写的货币资金说明。"


@pytest.mark.asyncio
async def test_inherit_skips_placeholder_prior_text(env):
    async with env["factory"]() as db:
        prior = await db.get(DisclosureNote, env["ids"]["prior_cash"])
        prior.text_content = "<p>本期无此项业务。</p>"
        await db.commit()
    resp = await env["client"].post(f"/api/projects/{env['ids']['p']}/notes/inherit-prior-year?year={YEAR}")
    assert resp.status_code == 200 and resp.json()["inherited"] == 0, resp.text
    cash = await _note(env, "cash")
    assert (cash.text_content, cash.updated_by) == (None, None)

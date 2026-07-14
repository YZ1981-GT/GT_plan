"""E1 导入导出字段、持久化及 SQL 契约测试。"""
from __future__ import annotations

import io
import json
import sys
from datetime import date, datetime
from pathlib import Path
from types import SimpleNamespace
from unittest.mock import AsyncMock, MagicMock
from uuid import UUID

import pytest
from fastapi import UploadFile
from hypothesis import given, settings, strategies as st
from openpyxl import Workbook, load_workbook

sys.path.insert(0, str(Path(__file__).resolve().parent.parent))

from app.routers.wp_render_strategies._e1_import_export import (
    _FIELD_MAPS,
    _FORMULA_FIELDS,
    _NUMERIC_FIELDS,
    _SHEET_HEADERS,
    _SHEET_ITEM_ID,
    _SUPPORTED_SHEETS,
    _create_template_wb,
    _export_row,
    _parse_row,
    _validate_columns,
    e1_import_data,
)

sheet_strategy = st.sampled_from(sorted(_SUPPORTED_SHEETS))

_CANONICAL_USER_FIELDS = {
    "E1-2": {"currency", "opening", "increase", "decrease", "fxRate", "adjustment", "note"},
    "E1-3": {"section", "group", "bankName", "totalLedgerBank", "accountNo", "accountType", "opening", "increase", "decrease", "adjustment", "statementBalance", "confirmAmount", "confirmIndexNo", "statementIndexNo", "reconciliationIndexNo", "restrictedAmount", "restrictedReason", "interestRate", "note", "fxCurrency", "fxRate", "openingFc", "increaseFc", "decreaseFc", "adjustmentFc"},
    "E1-4": {"seq", "bankName", "currency", "fxRate", "opening", "increase", "decrease", "adjustment", "queryBalance", "diffReason", "indexNo", "confirmationIndexNo", "note"},
    "E1-5": {"description", "category", "reportItem", "accountName", "noteItem", "debit", "credit", "indexNo", "note"},
    "E1-6": {"bankName", "accountNo", "bookBalance", "statementBalance", "bankReceivedCompanyUnreceivedItems", "bankPaidCompanyUnpaidItems", "companyReceivedBankUnreceivedItems", "companyPaidBankUnpaidItems", "diffReason"},
    "E1-7": {"denomination", "quantity"},
    "E1-8": {"currency", "denomination", "quantity", "fcAmount", "fxRate"},
    "E1-9": {"certNo", "bank", "depositor", "account", "certType", "currency", "depositDate", "maturityDate", "amount", "interestRate", "bookConsistent", "inconsistencyReason", "pledged", "pledgeMatter", "result", "certificateIndex", "openingProofIndex", "custodyProofIndex", "note"},
    "E1-10": {"bank", "accountNo", "accountType", "openDate", "accountStatus", "closeDate", "openReason", "closeReason", "companyInfoConsistent", "inconsistencyReason", "restrictionStatus", "openPurpose", "isNewThisPeriod", "isClosedThisPeriod", "hasBookRecord", "checkResult", "reason"},
    "E1-15": {"depositType", "bank", "accountNo", "annualRate", "balances", "bookInterests"},
    "E1-20": {"bank", "accountNo", "usage", "category", "currency", "fcAmount", "settleDate", "cutoffDate", "dailyRate", "fxRate", "note"},
    "E1-21": {"voucherNo", "date", "amount", "counterparty", "note"},
    "E1-22": {"voucherNo", "date", "amount", "counterparty", "note"},
}


_EXPECTED_ITEM_IDS = {
    "E1-2": "E1-cash-detail-rows", "E1-3": "E1-bank-detail-rows",
    "E1-4": "E1-digital-rows", "E1-5": "E1-adjustment-rows",
    "E1-6": "E1-reconciliation-rows", "E1-7": "E1-cash-count-rmb-rows",
    "E1-8": "E1-cash-count-fx-rows", "E1-9": "E1-cash-count-cert-rows",
    "E1-10": "E1-account-list-rows", "E1-15": "E1-interest-monthly-rows",
    "E1-20": "E1-accrued-interest-rows", "E1-21": "E1-cutoff-bank-rows",
    "E1-22": "E1-cutoff-other-rows",
}

_DATE_FIELDS = {
    "depositDate", "maturityDate", "openDate", "closeDate", "settleDate", "cutoffDate", "date",
}
_ENUM_VALUES = {
    ("E1-3", "section"): "accrued", ("E1-3", "group"): "institution",
    ("E1-5", "category"): "账项调整",
    ("E1-9", "bookConsistent"): "否", ("E1-9", "pledged"): "是",
    ("E1-9", "result"): "已见", ("E1-10", "companyInfoConsistent"): "不一致",
    ("E1-10", "isNewThisPeriod"): "Y", ("E1-10", "isClosedThisPeriod"): "N",
    ("E1-10", "hasBookRecord"): "Y", ("E1-10", "checkResult"): "一致",
    ("E1-20", "category"): "digital",
}


def _sample_row(sheet: str) -> dict:
    row = {"id": "source-row-id"}
    if sheet == "E1-6":
        item = {
            "id": "nested-1", "bankDate": "2026-01-02", "amount": 12.5,
            "postedAfterPeriod": True, "postingDate": "2026-01-05", "voucherNo": "记-1",
            "description": "未达账项", "counterAccount": "应收账款",
            "statementDate": "2026-01-03", "accountingCorrect": None,
            "adjustmentRequired": False, "note": "保留嵌套值",
        }
        row.update({
            "bankReceivedCompanyUnreceivedItems": [item],
            "bankPaidCompanyUnpaidItems": [{**item, "id": "nested-2", "amount": 2.5}],
            "companyReceivedBankUnreceivedItems": [{**item, "id": "nested-3", "amount": 8.5}],
            "companyPaidBankUnpaidItems": [{**item, "id": "nested-4", "amount": 1.5}],
        })
    elif sheet == "E1-15":
        row.update({
            "depositType": "七天通知存款", "bank": "矩阵银行", "accountNo": "A-001",
            "annualRate": 0.036, "balances": [month * 100.25 for month in range(1, 13)],
            "bookInterests": [month * 3.75 for month in range(1, 13)],
        })
    for field in _CANONICAL_USER_FIELDS[sheet]:
        if field in row:
            continue
        if (sheet, field) in _ENUM_VALUES:
            row[field] = _ENUM_VALUES[(sheet, field)]
        elif field in _DATE_FIELDS:
            row[field] = "2026-03-04"
        elif field in _NUMERIC_FIELDS[sheet]:
            row[field] = 3 if field == "seq" else 12.5
        else:
            row[field] = f"{field}-值"
    return row


def _round_trip(sheet: str, source: dict) -> dict:
    wb = _create_template_wb(sheet)
    wb.active.append(_export_row(sheet, source))
    buffer = io.BytesIO()
    wb.save(buffer)
    buffer.seek(0)
    ws = load_workbook(buffer, read_only=True, data_only=True).active
    headers = [str(cell.value).strip() for cell in next(ws.iter_rows(min_row=1, max_row=1)) if cell.value]
    values = next(ws.iter_rows(min_row=2, max_row=2, values_only=True))
    return _parse_row(sheet, values, headers)


@settings(max_examples=5)
@given(sheet=sheet_strategy)
def test_template_has_exact_configured_headers(sheet: str):
    ws = _create_template_wb(sheet).active
    actual = [str(cell.value).strip() for cell in next(ws.iter_rows(min_row=1, max_row=1)) if cell.value]
    assert actual == _SHEET_HEADERS[sheet]
    assert ws.freeze_panes == "A2"


def test_canonical_fields_and_item_ids_match_frontend_persistence_contract():
    assert _SHEET_ITEM_ID == _EXPECTED_ITEM_IDS
    assert {"E1-4", "E1-15"} <= _SUPPORTED_SHEETS
    for sheet, fields in _FIELD_MAPS.items():
        imported_fields = set(fields.values()) - _FORMULA_FIELDS[sheet]
        if sheet == "E1-15":
            assert imported_fields == {
                "depositType", "bank", "accountNo", "annualRate",
                *{f"balance{month}" for month in range(1, 13)},
                *{f"bookInterest{month}" for month in range(1, 13)},
            }
        else:
            assert imported_fields == _CANONICAL_USER_FIELDS[sheet]


@pytest.mark.parametrize("sheet", sorted(_SUPPORTED_SHEETS))
def test_field_values_round_trip_using_canonical_json(sheet: str):
    source = _sample_row(sheet)
    parsed = _round_trip(sheet, source)
    assert "id" in parsed and "rowId" not in parsed
    UUID(parsed["id"])
    for field in _CANONICAL_USER_FIELDS[sheet]:
        assert parsed[field] == source[field], f"{sheet}.{field} did not round-trip"
    assert not (_FORMULA_FIELDS[sheet] & parsed.keys())


@pytest.mark.parametrize("sheet", sorted(_SUPPORTED_SHEETS))
def test_formula_columns_are_exported_but_never_imported(sheet: str):
    formulas = _FORMULA_FIELDS[sheet]
    assert all(field in _FIELD_MAPS[sheet].values() for field in formulas)
    if not formulas:
        return
    source = _sample_row(sheet)
    source.update({field: 999999 for field in formulas})
    parsed = _round_trip(sheet, source)
    assert formulas.isdisjoint(parsed)


def test_export_recalculates_formula_values_from_canonical_inputs():
    sheet = "E1-20"
    source = {
        "bank": "测试银行", "accountNo": "001", "usage": "定期", "category": "bank",
        "currency": "人民币", "fcAmount": 1000, "settleDate": "2026-01-01",
        "cutoffDate": "2026-01-11", "dailyRate": 0.001, "fxRate": 1, "note": "",
    }
    values = dict(zip(_FIELD_MAPS[sheet].values(), _export_row(sheet, source)))
    assert values["days"] == 10
    assert values["accruedFc"] == pytest.approx(10)
    assert values["accruedRmb"] == pytest.approx(10)


def test_e1_3_layered_structure_and_professional_fields_round_trip():
    source = _sample_row("E1-3")
    source.update({
        "section": "accrued", "group": "finance", "statementIndexNo": "E1-3-S1",
        "reconciliationIndexNo": "E1-3-R1", "interestRate": 0.0185,
    })
    parsed = _round_trip("E1-3", source)
    assert parsed["section"] == "accrued"
    assert parsed["group"] == "finance"
    assert parsed["statementIndexNo"] == "E1-3-S1"
    assert parsed["reconciliationIndexNo"] == "E1-3-R1"
    assert parsed["interestRate"] == pytest.approx(0.0185)


def test_e1_4_diff_reason_is_user_field_and_round_trips_verbatim():
    source = _sample_row("E1-4")
    source["diffReason"] = "平台查询余额包含资产负债表日后入账交易，已取得流水核对。"
    parsed = _round_trip("E1-4", source)
    assert parsed["diffReason"] == source["diffReason"]
    assert "diffReason" not in _FORMULA_FIELDS["E1-4"]


def test_e1_15_account_matrix_exports_formulas_and_imports_two_twelve_month_arrays():
    source = _sample_row("E1-15")
    exported = dict(zip(_FIELD_MAPS["E1-15"].values(), _export_row("E1-15", source)))
    for month in range(1, 13):
        assert exported[f"balance{month}"] == source["balances"][month - 1]
        assert exported[f"bookInterest{month}"] == source["bookInterests"][month - 1]
        assert exported[f"calculatedInterest{month}"] == pytest.approx(
            source["balances"][month - 1] * source["annualRate"] / 12
        )
    parsed = _round_trip("E1-15", source)
    assert parsed["balances"] == source["balances"]
    assert parsed["bookInterests"] == source["bookInterests"]
    assert len(parsed["balances"]) == len(parsed["bookInterests"]) == 12
    assert not any(key.startswith("calculatedInterest") for key in parsed)


def test_e1_6_nested_outstanding_items_round_trip_without_sidecar_or_field_loss():
    source = _sample_row("E1-6")
    parsed = _round_trip("E1-6", source)
    for field in (
        "bankReceivedCompanyUnreceivedItems", "bankPaidCompanyUnpaidItems",
        "companyReceivedBankUnreceivedItems", "companyPaidBankUnpaidItems",
    ):
        assert parsed[field] == source[field]

    exported = dict(zip(_FIELD_MAPS["E1-6"].values(), _export_row("E1-6", source)))
    assert exported["companyReceived"] == pytest.approx(12.5)
    assert exported["companyPaid"] == pytest.approx(2.5)
    assert exported["bankReceived"] == pytest.approx(8.5)
    assert exported["bankPaid"] == pytest.approx(1.5)
    assert exported["reconciledBook"] == pytest.approx(22.5)
    assert exported["reconciledStatement"] == pytest.approx(19.5)
    assert exported["diff"] == pytest.approx(3.0)


def test_legacy_field_aliases_remain_export_compatible():
    values = dict(zip(
        _FIELD_MAPS["E1-5"].values(),
        _export_row("E1-5", {"debitAmount": 10, "creditAmount": 8, "indexRef": "A-1", "remark": "旧备注"}),
    ))
    assert values["debit"] == 10
    assert values["credit"] == 8
    assert values["indexNo"] == "A-1"
    assert values["note"] == "旧备注"


@pytest.mark.parametrize(
    ("sheet", "column", "raw", "field", "expected"),
    [
        ("E1-9", "存入日", datetime(2026, 2, 3, 15, 30), "depositDate", "2026-02-03"),
        ("E1-9", "到期日", date(2027, 4, 5), "maturityDate", "2027-04-05"),
        ("E1-10", "开户日期", "2026/06/07", "openDate", "2026-06-07"),
        ("E1-10", "本期新开", "是", "isNewThisPeriod", "Y"),
        ("E1-10", "本期注销", "no", "isClosedThisPeriod", "N"),
        ("E1-10", "清单核对一致", "否", "checkResult", "不一致"),
        ("E1-10", "企业信息核对一致", "否", "companyInfoConsistent", "不一致"),
        ("E1-9", "账面一致", "N", "bookConsistent", "否"),
        ("E1-9", "是否质押/受限", "Y", "pledged", "是"),
        ("E1-3", "一级区段", "（一）存款本金", "section", "principal"),
        ("E1-3", "一级区段", "应计利息", "section", "accrued"),
        ("E1-3", "二级分组", "存放财务公司款项", "group", "finance"),
        ("E1-3", "二级分组", "其他金融机构", "group", "institution"),
        ("E1-3", "二级分组", "principal", "group", "institution"),
        ("E1-20", "类别", "财务公司", "category", "finance"),
        ("E1-20", "类别", "银行", "category", "bank"),
        ("E1-20", "类别", "其他", "category", "other"),
        ("E1-20", "类别", "数字货币", "category", "digital"),
    ],
)
def test_date_and_enum_values_are_normalized(sheet: str, column: str, raw, field: str, expected: str):
    headers = _SHEET_HEADERS[sheet]
    row = [None] * len(headers)
    row[headers.index(column)] = raw
    parsed = _parse_row(sheet, tuple(row), headers)
    assert parsed[field] == expected


@settings(max_examples=5)
@given(sheet=sheet_strategy)
def test_column_tamper_is_rejected(sheet: str):
    wb = Workbook()
    ws = wb.active
    ws.append(["错误列A", "错误列B"])
    assert set(_validate_columns(ws, sheet)) == set(_SHEET_HEADERS[sheet])


class _Result:
    def __init__(self, row):
        self._row = row

    def fetchone(self):
        return self._row


@pytest.mark.asyncio
async def test_import_queries_project_before_upsert_and_writes_project_id_and_item_id():
    wb = _create_template_wb("E1-10")
    wb.active.append(_export_row("E1-10", {
        "bank": "甲银行", "accountNo": "001", "accountType": "基本户",
        "openDate": "2026/01/02", "openPurpose": "结算", "isNewThisPeriod": "是",
        "isClosedThisPeriod": "N", "hasBookRecord": "yes", "checkResult": "否", "reason": "待核实",
    }))
    content = io.BytesIO()
    wb.save(content)
    content.seek(0)

    db = AsyncMock()
    db.execute.side_effect = [
        _Result(SimpleNamespace(project_id="project-123")),
        _Result(None),
    ]
    upload = UploadFile(filename="e1.xlsx", file=content)
    response = await e1_import_data(
        wp_id="wp-123", sheet="E1-10", file=upload, db=db, current_user=MagicMock(),
    )

    assert response == {"row_count": 1, "sheet": "E1-10", "item_id": "E1-account-list-rows"}
    assert db.execute.await_count == 2
    select_call, insert_call = db.execute.await_args_list
    assert "SELECT project_id FROM working_paper" in str(select_call.args[0])
    assert select_call.args[1] == {"wp_id": "wp-123"}
    assert "project_id" in str(insert_call.args[0])
    params = insert_call.args[1]
    assert params["project_id"] == "project-123"
    assert params["wp_id"] == "wp-123"
    assert params["item_id"] == "E1-account-list-rows"
    stored = json.loads(params["remark"])
    assert stored[0]["id"] and "rowId" not in stored[0]
    assert stored[0]["openDate"] == "2026-01-02"
    assert stored[0]["isNewThisPeriod"] == "Y"
    assert stored[0]["hasBookRecord"] == "Y"
    assert stored[0]["checkResult"] == "不一致"
    db.commit.assert_awaited_once()


@pytest.mark.asyncio
async def test_import_does_not_upsert_when_working_paper_has_no_project():
    wb = _create_template_wb("E1-15")
    wb.active.append(_export_row("E1-15", {
        "depositType": "活期存款", "bank": "甲银行", "accountNo": "001",
        "annualRate": 0.012, "balances": [100] * 12, "bookInterests": [1] * 12,
    }))
    content = io.BytesIO()
    wb.save(content)
    content.seek(0)
    db = AsyncMock()
    db.execute.return_value = _Result(None)

    with pytest.raises(Exception) as exc_info:
        await e1_import_data(
            wp_id="missing", sheet="E1-15",
            file=UploadFile(filename="e1.xlsx", file=content), db=db, current_user=MagicMock(),
        )
    assert getattr(exc_info.value, "status_code", None) == 404
    assert db.execute.await_count == 1
    db.commit.assert_not_awaited()

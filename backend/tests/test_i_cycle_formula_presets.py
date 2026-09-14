"""I 类公式预设合法性契约测试。

验证 `prefill_formula_mapping.json` 中 I 类（wp_code 以 I 开头且后续全为数字）
所有块的：

1. sheet 存在性（源模板 xlsx sheetnames）
2. 科目存在性（standard_account_chart.json）
3. 防成环（审定表 WP() 不引用自己；明细表不引用回审定表）
4. wp_name 语义
5. 无 1712/1717/1911 残留

spec: .kiro/specs/i-cycle-four-table-extraction-and-disclosure-alignment/
"""
from __future__ import annotations

import json
import re
from pathlib import Path
from typing import Any

import pytest

# ── 路径 ──────────────────────────────────────────────────────────────────────

ROOT = Path(__file__).resolve().parents[1]
DATA_DIR = ROOT / "data"
TEMPLATE_DIR = ROOT / "wp_templates" / "I"
MAPPING_PATH = DATA_DIR / "prefill_formula_mapping.json"
ACCOUNT_CHART_PATH = DATA_DIR / "standard_account_chart.json"


# ── 数据加载 ──────────────────────────────────────────────────────────────────

def _load_mapping() -> list[dict[str, Any]]:
    raw = json.loads(MAPPING_PATH.read_text(encoding="utf-8"))
    mappings = raw["mappings"] if isinstance(raw, dict) and "mappings" in raw else raw
    return [
        b for b in mappings
        if b.get("wp_code", "").startswith("I")
        and b["wp_code"][1:].isdigit()
        and not b.get("deleted", False)
    ]


def _load_account_codes() -> set[str]:
    raw = json.loads(ACCOUNT_CHART_PATH.read_text(encoding="utf-8"))
    return {a["code"] for a in raw.get("accounts", [])}


def _load_template_sheetnames() -> dict[str, list[str]]:
    """返回 {wp_code: [sheetname, ...]}，跳过 ~$ 锁文件。"""
    import openpyxl

    result: dict[str, list[str]] = {}
    for xlsx_path in sorted(TEMPLATE_DIR.glob("*.xlsx")):
        if xlsx_path.name.startswith("~$"):
            continue
        # 提取 wp_code: 文件名形如 "I1 无形资产…xlsx"
        stem = xlsx_path.stem
        wp_code_match = re.match(r"^(I\d+)\s", stem)
        if not wp_code_match:
            continue
        wp_code = wp_code_match.group(1)
        wb = openpyxl.load_workbook(xlsx_path, read_only=True)
        result[wp_code] = list(wb.sheetnames)
        wb.close()
    return result


# ── Fixtures ──────────────────────────────────────────────────────────────────

@pytest.fixture(scope="module")
def i_blocks() -> list[dict[str, Any]]:
    return _load_mapping()


@pytest.fixture(scope="module")
def account_codes() -> set[str]:
    return _load_account_codes()


@pytest.fixture(scope="module")
def template_sheets() -> dict[str, list[str]]:
    return _load_template_sheetnames()


# ── 参数化 ────────────────────────────────────────────────────────────────────

def _block_ids() -> list[str]:
    """为每个块生成唯一 test id。"""
    blocks = _load_mapping()
    return [f"{b['wp_code']}-{b['sheet']}" for b in blocks]


def _blocks_parametrize():
    blocks = _load_mapping()
    return blocks


# ── 1. Sheet 存在性 ──────────────────────────────────────────────────────────

#: 曾经「贴错标签」的 sheet 名 → 源 xlsx 真实 tab 名（**历史台账，非逃逸阀**）。
#:
#: 🔴 2026-08-09（Task 7/9）：改造前这里是一张 `pytest.skip` 白名单 ——
#: ``("I1", "审定表I1-1"): "审定表I1"`` 把「块的 sheet 字段指向源模板不存在的 tab」
#: 这一真缺陷**放行**了，于是该守卫长期 `1 skipped` 而缺陷留在数据里。
#: Task 7 已把预设块的 `sheet` 字段改成 `审定表I1`，故该条目**已过期必须移除**
#: （memory 铁律「已修好必删」）。
#:
#: 本表保留下来只作两件事：
#:   1. 记载历史（下个会话看到 `审定表I1` 时知道它曾是 `审定表I1-1`）；
#:   2. 由 :func:`test_no_stale_sheet_label_whitelist` 做 **stale 检测** ——
#:      表里的错名一旦在数据里重新出现即打红，而不是被 skip 掉。
_HISTORICAL_SHEET_LABEL_FIXES: dict[tuple[str, str], str] = {
    ("I1", "审定表I1-1"): "审定表I1",
}


@pytest.mark.parametrize("block", _blocks_parametrize(), ids=_block_ids())
def test_sheet_exists_in_template(block: dict, template_sheets: dict[str, list[str]]) -> None:
    """每个块的 sheet 字段必须存在于对应源模板 xlsx 的 sheetnames 中。

    🔴 **不再有白名单** —— 源模板是唯一裁决者，预设必须对齐它（R3.1/R3.6）。
    """
    wp_code = block["wp_code"]
    sheet = block["sheet"]

    assert wp_code in template_sheets, (
        f"源模板目录下找不到 {wp_code} 对应的 xlsx 文件"
    )
    sheetnames = template_sheets[wp_code]
    hint = ""
    if (wp_code, sheet) in _HISTORICAL_SHEET_LABEL_FIXES:
        hint = (
            f"\n🔴 该 sheet 名曾于 2026-08-09 被修正为 "
            f"'{_HISTORICAL_SHEET_LABEL_FIXES[(wp_code, sheet)]}'，"
            "本次出现说明数据被回退（`prefill_formula_mapping.json` 是多 spec 共享的"
            "回退高发文件）⇒ 重跑 "
            "`python backend/scripts/fix/fix_i_cycle_prefill_presets.py --apply`"
        )
    assert sheet in sheetnames, (
        f"{wp_code} 的 sheet '{sheet}' 不存在于模板中。"
        f" 可用的 sheet: {sheetnames}{hint}"
    )


def test_no_stale_sheet_label_whitelist(template_sheets: dict[str, list[str]]) -> None:
    """历史台账里的每个「错名」都必须**确实是错的**（stale 检测）。

    防两种退化：

    * 台账变成逃逸阀 —— 有人把真实存在的 tab 名塞进来当豁免；
    * 台账过期 —— 源模板某天真的加了 `审定表I1-1` 这个 tab，此时台账的
      「这是错名」前提不再成立，应移除该条目而不是继续记载。
    """
    for (wp_code, wrong), right in _HISTORICAL_SHEET_LABEL_FIXES.items():
        sheetnames = template_sheets.get(wp_code, [])
        assert wrong not in sheetnames, (
            f"台账过期：'{wrong}' 现在**确实存在**于 {wp_code} 源模板中"
            f"（{sheetnames}）⇒ 它不再是错名，请从 "
            "_HISTORICAL_SHEET_LABEL_FIXES 移除该条目"
        )
        assert right in sheetnames, (
            f"台账的「正确名」'{right}' 不在 {wp_code} 源模板中（{sheetnames}）"
            " ⇒ 台账本身写错了"
        )


# ── 2. 科目存在性 ────────────────────────────────────────────────────────────

# 已知 chart_conflict：I2 的 1703 在本项目被诊断冲突后回退兜底 1704，
# 但 1704 不在 standard_account_chart.json（CAS 标准没有该码）。
# 见 memory §任务状态 "I2 的 `1703` 被 `chart_conflict` 正确诊断并回退兜底 `1704`"
_KNOWN_CHART_CONFLICTS: set[str] = {"1704"}


@pytest.mark.parametrize("block", _blocks_parametrize(), ids=_block_ids())
def test_account_codes_exist_in_chart(block: dict, account_codes: set[str]) -> None:
    """每个块 account_codes 里的码必须存在于标准科目表。
    空列表合法（I5）。已知 chart_conflict 码走白名单。
    """
    codes = block.get("account_codes", [])
    for code in codes:
        if code in _KNOWN_CHART_CONFLICTS:
            continue
        assert code in account_codes, (
            f"{block['wp_code']} / {block['sheet']} 的科目码 '{code}' "
            f"不在 standard_account_chart.json 中"
        )


# ── 3. 防成环 ────────────────────────────────────────────────────────────────

_WP_PATTERN = re.compile(r"WP\(\s*'([^']+)'\s*,\s*'([^']+)'")


@pytest.mark.parametrize("block", _blocks_parametrize(), ids=_block_ids())
def test_no_self_referencing_cycles(block: dict) -> None:
    """
    审定表块的 cells[].formula 里的 WP() 引用不得指向自己
    （如 WP('I1','审定表I1',…) 不应出现在 I1 审定表块里）；
    明细表块不得有 WP() 引用回审定表。
    """
    wp_code = block["wp_code"]
    sheet = block["sheet"]
    cells = block.get("cells", [])

    is_adjudication = "审定表" in sheet
    is_detail = "明细表" in sheet

    for cell in cells:
        formula = cell.get("formula", "")
        for match in _WP_PATTERN.finditer(formula):
            ref_wp = match.group(1)
            ref_sheet = match.group(2)

            if is_adjudication:
                # 审定表不得引用自己的审定表 sheet
                if ref_wp == wp_code and "审定表" in ref_sheet:
                    pytest.fail(
                        f"{wp_code}/{sheet} 审定表块 WP() 引用了自己的审定表: "
                        f"WP('{ref_wp}','{ref_sheet}',...) "
                        f"cell_ref={cell.get('cell_ref')}"
                    )

            if is_detail:
                # 明细表不得引用回审定表
                if ref_wp == wp_code and "审定表" in ref_sheet:
                    pytest.fail(
                        f"{wp_code}/{sheet} 明细表块 WP() 引用了审定表: "
                        f"WP('{ref_wp}','{ref_sheet}',...) "
                        f"cell_ref={cell.get('cell_ref')}"
                    )


# ── 4. wp_name 语义 ──────────────────────────────────────────────────────────

_WP_NAME_RULES: dict[str, list[str]] = {
    "I1": ["无形资产", "摊销"],
    "I2": ["开发支出", "研发"],
    "I3": ["商誉"],
    "I4": ["长期待摊费用", "待摊", "摊销"],
    "I5": ["其他非流动资产", "非流动"],
    "I6": ["研发费用", "研发"],
}


@pytest.mark.parametrize("block", _blocks_parametrize(), ids=_block_ids())
def test_wp_name_semantic(block: dict) -> None:
    """wp_name 必须包含对应循环的语义关键字。"""
    wp_code = block["wp_code"]
    wp_name = block.get("wp_name", "")
    keywords = _WP_NAME_RULES.get(wp_code, [])

    if not keywords:
        return  # 没有定义规则的跳过

    matched = any(kw in wp_name for kw in keywords)
    assert matched, (
        f"{wp_code} 块 wp_name='{wp_name}' 不含任何预期关键字 {keywords}"
    )


# ── 5. 无 1712/1717/1911 残留 ────────────────────────────────────────────────

_FORBIDDEN_CODES = {"1712", "1717", "1911"}


@pytest.mark.parametrize("block", _blocks_parametrize(), ids=_block_ids())
def test_no_forbidden_code_residue(block: dict) -> None:
    """I 类所有块的 account_codes 与 cells[].formula 中不得出现 1712/1717/1911。"""
    wp_code = block["wp_code"]
    sheet = block["sheet"]

    # 检查 account_codes
    codes = block.get("account_codes", [])
    for code in codes:
        assert code not in _FORBIDDEN_CODES, (
            f"{wp_code}/{sheet} account_codes 含禁用码 '{code}'"
        )

    # 检查 cells[].formula
    for cell in block.get("cells", []):
        formula = cell.get("formula", "")
        for forbidden in _FORBIDDEN_CODES:
            assert forbidden not in formula, (
                f"{wp_code}/{sheet} cell_ref='{cell.get('cell_ref')}' "
                f"公式含禁用码 '{forbidden}': {formula}"
            )


# ─────────────────────────────────────────────────────────────────────────────
# Property 10~13（spec `i-cycle-extraction-formula-and-disclosure-closure` Task 9）
#
# 🔴 在**本文件内**扩充，禁另建同域守卫文件 —— 同一不变式两处各写一份，
#    改一处另一处不红（memory 已记的判据双真源族）。
# ─────────────────────────────────────────────────────────────────────────────

import hashlib  # noqa: E402
import os  # noqa: E402
import subprocess  # noqa: E402
import sys  # noqa: E402

#: 幂等脚本路径（唯一写者，禁另建同域脚本）
_FIX_SCRIPT = ROOT / "scripts" / "fix" / "fix_i_cycle_prefill_presets.py"


def _run_fix(*args: str) -> subprocess.CompletedProcess:
    """跑幂等脚本。

    🔴 必须显式 ``encoding='utf-8'`` —— Windows 下 ``text=True`` 用 locale(GBK)
    解码含中文的 stdout 会抛 UnicodeDecodeError 使 ``stdout`` 变 ``None``，
    断言随后以 ``TypeError`` 失败 = 既判不出欠账也判不出脚本是好的
    （memory 已记：backend 还有约 30 处同款）。
    """
    env = dict(os.environ)
    env["PYTHONIOENCODING"] = "utf-8"
    env["PYTHONUTF8"] = "1"
    return subprocess.run(
        [sys.executable, str(_FIX_SCRIPT), *args],
        cwd=str(ROOT.parent),
        env=env,
        capture_output=True,
        encoding="utf-8",
        errors="replace",
        timeout=300,
    )


class TestProperty10Idempotent:
    """Property 10：``--check`` 归零 + 二次 ``--apply`` 后 md5 逐字节不变。"""

    def test_check_returns_zero(self):
        r = _run_fix("--check")
        assert r.stdout is not None, "stdout 为 None（编码参数缺失，见 _run_fix docstring）"
        assert r.returncode == 0, (
            f"`--check` 未归零（rc={r.returncode}）：\n{r.stdout}\n{r.stderr}"
        )

    def test_second_apply_is_byte_identical(self):
        """二次 apply 必须 md5 不变、rc=0、且 stderr 无 traceback。

        🔴 三条缺一不可 —— 脚本崩在写盘**之前**时 md5 天然不变，
        只看 md5 会把 ``NameError`` 判成「幂等」（本轮已实证一次）。
        """
        before = hashlib.md5(MAPPING_PATH.read_bytes()).hexdigest()
        r = _run_fix("--apply")
        after = hashlib.md5(MAPPING_PATH.read_bytes()).hexdigest()

        assert r.returncode == 0, f"二次 apply rc={r.returncode}：\n{r.stderr}"
        assert "Traceback" not in (r.stderr or ""), (
            f"二次 apply 有 traceback（崩在写盘前 → md5 天然不变）：\n{r.stderr}"
        )
        assert before == after, "二次 apply 改写了文件（非幂等）"
        assert "无需修改" in (r.stdout or ""), (
            f"二次 apply 仍报变更（同一字段可能有多个写者）：\n{r.stdout}"
        )

    def test_round_trip_reproduces_source(self):
        """round-trip：``json.dumps`` 能逐字节复现原文。

        防「脚本重排整个 1 MB JSON」→ 与并发会话互相回退
        （该文件被 D/G/H/I/K/L/N 多 spec 共享）。
        """
        raw = MAPPING_PATH.read_text(encoding="utf-8")
        rt = json.dumps(json.loads(raw), ensure_ascii=False, indent=2) + "\n"
        assert rt == raw, "json.dumps 无法逐字节复现原文（缩进/尾换行不符）"


class TestProperty11PrevSheetArgSynced:
    """Property 11：改块 sheet 名时，块内公式实参里的旧 sheet 名必须同步改。

    🔴 改造前的缺陷形态 = 只改公式实参、**不改块自己的 `sheet` 字段**
    （脚本还把错名当查找键），于是块永久指向源模板不存在的 tab
    而 ``--check`` 仍 rc=0。
    """

    _SHEET_ARG = re.compile(r"(?:PREV|WP)\(\s*'[^']+'\s*,\s*'([^']+)'")

    def test_no_block_sheet_is_absent_from_template(self, template_sheets):
        """每个块的 `sheet` 字段必须是源模板真实 tab（与 Property 8 同源，此处防回退）。"""
        for b in _load_mapping():
            wp, sheet = b["wp_code"], b["sheet"]
            names = template_sheets.get(wp)
            if not names:
                continue
            assert sheet in names, (
                f"{wp} 块 sheet='{sheet}' 不在源模板 tab 名集合内。"
                f"\n可用: {names}"
            )

    def test_formula_sheet_args_are_real_tabs(self, template_sheets):
        """公式实参里的 sheet 名（同 wp_code 内）必须是源模板真实 tab。"""
        bad: list[str] = []
        for b in _load_mapping():
            wp = b["wp_code"]
            names = template_sheets.get(wp) or []
            for c in b.get("cells", []):
                f = c.get("formula", "")
                for m in self._SHEET_ARG.finditer(f):
                    arg = m.group(1)
                    # 只校验「指向本 wp_code 自己 sheet」的实参
                    if not arg.startswith(("审定表", "明细表", "摊销测算")):
                        continue
                    if wp not in arg:
                        continue  # 跨循环引用，交由对方循环的守卫管
                    if names and arg not in names:
                        bad.append(
                            f"{wp}/{b['sheet']} cell='{c.get('cell_ref')}' "
                            f"公式实参 sheet '{arg}' 不存在于源模板"
                        )
        assert not bad, "公式实参 sheet 名与源模板不符：\n" + "\n".join(bad)

    def test_reverse_self_check_pattern_matches(self):
        """反向自检：正则确实能从 PREV/WP 里抽出 sheet 实参（防判据空转）。"""
        got = self._SHEET_ARG.findall("=PREV('I1','审定表I1','未审数')")
        assert got == ["审定表I1"], f"实参抽取失效：{got}"


class TestProperty12DisclosureRegistryComplete:
    """Property 12：12 张披露 sheet 必须或有预设块、或在显式无预设登记表内。"""

    def _registry(self):
        import importlib.util

        spec = importlib.util.spec_from_file_location("_fix_i", _FIX_SCRIPT)
        assert spec and spec.loader
        mod = importlib.util.module_from_spec(spec)
        sys.modules["_fix_i"] = mod  # 🔴 dataclass 解析需要它在 sys.modules 里
        spec.loader.exec_module(mod)
        return mod

    def test_every_disclosure_sheet_has_disposition(self, template_sheets):
        mod = self._registry()
        declared = {(b["wp_code"], b["sheet"]) for b in _load_mapping()}
        missing = [
            k
            for k in mod.DISCLOSURE_SHEETS
            if k not in declared and k not in mod.DISCLOSURE_NO_PRESET_REGISTRY
        ]
        assert not missing, (
            f"披露 sheet 既无预设块也未登记（R4.1 禁沉默）：{missing}"
        )

    def test_registry_sheets_are_real_tabs(self, template_sheets):
        """登记表里的 sheet 名必须是源模板真实 tab（防登记一个不存在的名字）。"""
        mod = self._registry()
        bad = [
            (wp, sh)
            for wp, sh in mod.DISCLOSURE_NO_PRESET_REGISTRY
            if sh not in (template_sheets.get(wp) or [])
        ]
        assert not bad, f"登记表 sheet 名不在源模板 tab 集合内：{bad}"

    def test_registry_count_anchor(self):
        """条目数锚点（6 循环 × 2 变体），防悄悄删条目让判据失去覆盖。"""
        mod = self._registry()
        assert len(mod.DISCLOSURE_SHEETS) == 12
        assert len(mod.DISCLOSURE_NO_PRESET_REGISTRY) == 12

    def test_registry_reasons_have_evidence(self):
        """每条理由必须够长且含实证标记（防「暂不需要」这类空话登记）。"""
        mod = self._registry()
        thin = [
            (k, len(v))
            for k, v in mod.DISCLOSURE_NO_PRESET_REGISTRY.items()
            if len(v) < 30
        ]
        assert not thin, f"登记理由过短（须写明实证依据）：{thin}"


class TestProperty13NoHardcodedClientCode:
    """Property 13：披露块公式不得出现项目专属硬编码（如 `AUX('1701','类别','具体编码')`）。

    当前 12 张披露 sheet 全在无预设登记表内 ⇒ 本判据是**防回退**性质：
    将来若有人给披露 sheet 加预设块，硬编码客户码会立刻打红。
    """

    #: 形如 AUX('1701','客户','007960') —— 第三参是具体编码/名称
    _AUX3 = re.compile(r"AUX\(\s*'[^']+'\s*,\s*'[^']+'\s*,\s*'[^']+'")

    def test_no_project_specific_aux_in_any_block(self):
        bad: list[str] = []
        for b in _load_mapping():
            for c in b.get("cells", []):
                f = c.get("formula", "")
                if self._AUX3.search(f):
                    bad.append(
                        f"{b['wp_code']}/{b['sheet']} cell='{c.get('cell_ref')}': {f}"
                    )
        assert not bad, (
            "预设含项目专属硬编码（AUX 三参形态，换项目即失效）：\n" + "\n".join(bad)
        )

    def test_reverse_self_check_pattern_matches(self):
        """反向自检：该正则确实能命中三参 AUX、且不误伤两参 AUX。"""
        assert self._AUX3.search("=AUX('1511.01','客户','007960','期末余额')")
        assert not self._AUX3.search("=AUX('1701','类别')")


class TestProperty14CrossCycleAccountOwnership:
    """Property 14：预设块引用的科目码必须属于**本循环**，跨循环引用须显式登记。

    🔴 为什么「科目存在性」判据（本文件第 2 组）不足以覆盖这一条：
    ``6602 管理费用`` 确实存在于 ``standard_account_chart.json``（K9 的科目），
    故「用了别的循环的科目」这类错**完全逃过存在性检查**。改造前 I2 明细表的
    ``研发费用_期末`` 就写着 ``TB('6602')``（研发费用真源是 ``6604``），
    而该文件当时 71 passed 全绿 —— 变异检验（M6）才把这个缺口暴露出来。

    判据 = 每块的 ``formula`` 抽出码 ∪ ``account_codes`` 必须 ⊆
    「本循环 :data:`I_CYCLE_SEGMENTS` 各段 ``fallback`` 的一级码集合」，
    唯一例外是 :data:`_CROSS_CYCLE_TIE_OUTS` 显式登记的跨循环勾稽。
    """

    #: 单码正则。**必须先把区间 `SUM_TB('a~b')` 整体消费掉再抽单码** ——
    #: 否则区间上界会被当成「本循环引用了别的科目」误报（memory 已记的 H1 同款坑）。
    _RANGE = re.compile(r"(?:SUM_TB|TB_SUM)\(\s*'(\d+)\s*~\s*(\d+)'")
    _CODE = re.compile(r"'(\d{4})(?:[.\-]\d+)*'")

    #: 合法的**跨循环勾稽**登记表：``(wp_code, sheet, 外来码) -> (归属循环, 理由)``。
    #:
    #: 🔴 三条约束（由 :meth:`test_registered_tie_outs_are_well_formed` 钉死）：
    #:   1. 外来码必须确实是**另一个 I 循环**声明的兜底码（不是随便一个存在的码）——
    #:      否则「登记」就成了万能逃逸阀；
    #:   2. 理由须写明勾稽方向与源模板依据（≥20 字）；
    #:   3. **stale 检测**：登记项若在数据里已不复现即打红（防条目永久留存变成盲区）。
    _CROSS_CYCLE_TIE_OUTS: dict[tuple[str, str, str], tuple[str, str]] = {
        ("I2", "明细表I2-2", "6604"): (
            "I6",
            "开发支出明细表需与 I6 研发费用勾稽（资本化 1704 + 费用化 6604 = 研发投入总额）；"
            "cell_ref='研发费用_期末' 的 description 已写明该用途。"
            "🔴 改造前此处误写 6602（管理费用，归 K9），Task 7 已修正。",
        ),
    }

    def _segments_mod(self):
        from app.services.four_table import i_cycle_accounts as mod

        return mod

    def _own_codes(self, wp_code: str) -> set[str]:
        """本循环声明的一级兜底码集合（各段 ``fallback`` 并集，取前 4 位）。"""
        mod = self._segments_mod()
        out: set[str] = set()
        for seg in mod.I_CYCLE_SEGMENTS.get(wp_code, ()):
            for code in seg.fallback:
                out.add(str(code)[:4])
        return out

    def _codes_in_formula(self, formula: str) -> set[str]:
        """抽出公式引用的一级科目码（**区间端点不算引用**）。"""
        stripped = self._RANGE.sub("", formula or "")
        return {m.group(1) for m in self._CODE.finditer(stripped)}

    def test_all_referenced_codes_belong_to_own_cycle(self):
        """核心判据：外来码必须在登记表内。"""
        bad: list[str] = []
        for b in _load_mapping():
            wp, sheet = b["wp_code"], b["sheet"]
            allowed = self._own_codes(wp)
            if not allowed:
                continue  # I5 兜底码为空 tuple（宁缺勿造），无判据基础

            used: set[str] = set()
            for c in b.get("cells", []):
                used |= self._codes_in_formula(c.get("formula", ""))
            used |= {str(x)[:4] for x in b.get("account_codes", [])}

            for code in sorted(used - allowed):
                if (wp, sheet, code) in self._CROSS_CYCLE_TIE_OUTS:
                    continue
                bad.append(
                    f"{wp}/{sheet} 引用了非本循环科目 '{code}'"
                    f"（本循环声明码 {sorted(allowed)}）。"
                    f"若属跨循环勾稽须登记进 _CROSS_CYCLE_TIE_OUTS 并写明依据"
                )
        assert not bad, (
            "预设引用了别的循环的科目（科目存在性检查抓不到这类错）：\n" + "\n".join(bad)
        )

    def test_registered_tie_outs_are_well_formed(self):
        """登记项三向自检：归属真实 + 理由实质 + stale 检测。"""
        mod = self._segments_mod()
        all_i_codes = {
            wp: self._own_codes(wp) for wp in mod.I_CYCLE_SEGMENTS
        }
        blocks = {(b["wp_code"], b["sheet"]): b for b in _load_mapping()}

        problems: list[str] = []
        for (wp, sheet, code), (owner, reason) in self._CROSS_CYCLE_TIE_OUTS.items():
            # 1. 外来码必须真属于登记的那个 I 循环
            if code not in all_i_codes.get(owner, set()):
                problems.append(
                    f"({wp},{sheet},{code}) 登记归属 {owner}，但 {owner} 的声明码是"
                    f" {sorted(all_i_codes.get(owner, set()))} —— 登记不成立"
                )
            # 2. 理由须实质
            if len(reason) < 20:
                problems.append(f"({wp},{sheet},{code}) 理由过短（{len(reason)} 字）")
            # 3. stale 检测：该块必须仍在、且确实仍引用该码
            b = blocks.get((wp, sheet))
            if b is None:
                problems.append(f"({wp},{sheet},{code}) 块已不存在 —— 登记项已过期须移除")
                continue
            used: set[str] = {str(x)[:4] for x in b.get("account_codes", [])}
            for c in b.get("cells", []):
                used |= self._codes_in_formula(c.get("formula", ""))
            if code not in used:
                problems.append(
                    f"({wp},{sheet},{code}) 数据里已不再引用该码 —— 登记项已过期须移除"
                )
        assert not problems, "跨循环勾稽登记表自检失败：\n" + "\n".join(problems)

    def test_reverse_self_check_range_endpoints_not_counted(self):
        """反向自检：区间端点不得被当成引用（防误报本循环外科目）。"""
        got = self._codes_in_formula("=SUM_TB('1401~1499','期末余额')+TB('1704','期末余额')")
        assert got == {"1704"}, f"区间端点未被排除：{got}"

    def test_reverse_self_check_detects_foreign_code(self):
        """反向自检：判据确实能识别外来码（防判据空转）。"""
        allowed = self._own_codes("I2")
        assert allowed, "I2 声明码为空，判据无基础"
        foreign = self._codes_in_formula("=TB('6602','期末余额')") - allowed
        assert foreign == {"6602"}, (
            f"判据无法识别 6602（管理费用，归 K9）为外来码：{foreign}"
        )

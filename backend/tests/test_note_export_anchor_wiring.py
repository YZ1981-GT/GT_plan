"""附注导出路径锚点写入接线测试 — deliverable-lineage-wiring-and-writeback-closure Task 2.1

**关键**：本测试走 `NoteWordExporter._export_template_mode_with_meta` 的**生产代码路径**
（只把模板文件与 DB 读取替身化），断言锚点确实被写入、meta 确实带出 —— 这是防
「前序 spec 22/22 全绿而生产从未调用 write_section_anchors」那类假绿的关键手段：
只测 `write_section_anchors` 纯函数本身永远发现不了「没人调用它」。

覆盖：
- Property 1：锚点集合 == kept_codes（裁剪章节不写）
- Property 2：写锚点前后可见段落文字序列逐字相等
- meta 三项（kept_codes / anchor_map / rendered_block_hashes）与文档实际一致
- 反向自检：跳过锚点写入步骤则 scan_anchor_blocks 为空（证明该步骤不可省）
"""

from __future__ import annotations

from dataclasses import dataclass
from io import BytesIO
from types import SimpleNamespace
from uuid import uuid4

import pytest
from docx import Document

from app.services import note_word_exporter as nwe
from app.services.note_word_exporter import NoteWordExporter
from app.services.section_anchor_utils import (
    anchor_name,
    block_text_hash,
    scan_anchor_blocks,
)

SOE_CODES = ["八、1", "八、2", "八、3"]


@dataclass
class _FakeEntry:
    abs_path: str
    exists: bool = True


def _make_template(tmp_path, codes: list[str]) -> str:
    """构造含 ``##SECTION:code##`` 块的合成模板 docx（不含填充占位符）。"""
    doc = Document()
    doc.add_paragraph("附注模板首段")
    for code in codes:
        doc.add_paragraph(f"##SECTION:{code}##")
        doc.add_paragraph(f"{code} 章节正文一")
        doc.add_paragraph(f"{code} 章节正文二")
        doc.add_paragraph(f"##/SECTION:{code}##")
    path = tmp_path / "note_template.docx"
    doc.save(str(path))
    return str(path)


def _fake_note(code: str):
    return SimpleNamespace(
        note_section=code,
        section_title=f"{code} 标题",
        text_content=f"{code} DB 文字",
        table_data={},
        is_deleted=False,
    )


@pytest.fixture()
def wired(monkeypatch, tmp_path):
    """把模板解析与 DB 读取替身化，其余全部走生产代码。"""
    template_path = _make_template(tmp_path, SOE_CODES)

    monkeypatch.setattr(
        nwe,
        "get_template_manifest_loader",
        lambda: SimpleNamespace(
            resolve_disclosure_notes=lambda tt, rs: _FakeEntry(template_path)
        ),
    )
    monkeypatch.setattr(nwe, "_load_section_code_index", lambda vk: [])
    monkeypatch.setattr(nwe, "note_applies_to_report_scope", lambda c, tt, rs: True)
    monkeypatch.setattr(nwe, "should_skip_empty_section", lambda d: False)
    # 内容控件与锚点是两条独立机制，本测试聚焦锚点 → 显式关闭控件（不依赖默认值）
    monkeypatch.setattr(
        nwe.settings, "DELIVERABLE_LINEAGE_CONTENT_CONTROL_ENABLED", False
    )
    return template_path


async def _run_export(exporter: NoteWordExporter):
    return await exporter._export_template_mode_with_meta(
        uuid4(),
        2025,
        template_type="soe",
        report_scope="standalone",
        sections=None,
    )


def _visible_paragraphs(buf: BytesIO) -> list[str]:
    """按文档顺序取全部段落文字。

    **必须下钻 w:sdt** —— `Document.paragraphs` 只返回 body 级 `w:p`，内容控件开启后
    章节内容被包进 `w:sdt/w:sdtContent`，用 `doc.paragraphs` 会「看不见」它们，从而把
    「注入改变了可见文字」误判出来（本测试首版即因此假红）。
    """
    from docx.oxml.ns import qn

    buf.seek(0)
    doc = Document(buf)
    ns_w = "http://schemas.openxmlformats.org/wordprocessingml/2006/main"
    out: list[str] = []
    for p in doc.element.body.iter(qn("w:p")):
        out.append("".join(t.text or "" for t in p.iter(f"{{{ns_w}}}t")))
    return out


# ─── Property 1 + meta 一致性 ────────────────────────────────────────────────


@pytest.mark.asyncio
async def test_property_1_anchors_cover_exactly_kept_sections(wired, monkeypatch):
    # Feature: deliverable-lineage-wiring-and-writeback-closure, Property 1
    """裁剪掉 八、2 后：锚点集合 == 剩余保留章节，被裁剪章节无锚点。"""
    exporter = NoteWordExporter(db=None)

    async def _load_notes(self, pid, year, sections):
        # 八、2 内容为空 → should_skip_empty_section 判 True → delete_section_block
        empty = _fake_note("八、2")
        empty.text_content = ""
        empty.table_data = {}
        return [_fake_note("八、1"), empty, _fake_note("八、3")]

    monkeypatch.setattr(NoteWordExporter, "_load_notes", _load_notes)
    monkeypatch.setattr(
        nwe,
        "should_skip_empty_section",
        lambda d: not (d or {}).get("text_content") and not (d or {}).get("table_data"),
    )

    buf, meta = await _run_export(exporter)

    assert set(meta.kept_codes) == {"八、1", "八、3"}
    assert set(meta.anchor_map) == {"八、1", "八、3"}
    assert meta.anchor_map["八、1"] == anchor_name("八、1")

    buf.seek(0)
    blocks = scan_anchor_blocks(Document(buf))
    assert {b.section_code for b in blocks} == {"八、1", "八、3"}
    assert anchor_name("八、2") not in {b.anchor_name for b in blocks}


@pytest.mark.asyncio
async def test_meta_rendered_hashes_match_document(wired, monkeypatch):
    """meta.rendered_block_hashes 与交付文档中各章节块实际内容哈希一致。"""
    exporter = NoteWordExporter(db=None)

    async def _load_notes(self, pid, year, sections):
        return [_fake_note(c) for c in SOE_CODES]

    monkeypatch.setattr(NoteWordExporter, "_load_notes", _load_notes)

    buf, meta = await _run_export(exporter)
    assert set(meta.rendered_block_hashes) == set(SOE_CODES)

    buf.seek(0)
    for block in scan_anchor_blocks(Document(buf)):
        assert (
            block_text_hash(block.elements)
            == meta.rendered_block_hashes[block.section_code]
        ), f"{block.section_code} 的 Rendered_Block_Hash 与交付文档实际内容不一致"


@pytest.mark.asyncio
async def test_anchor_survives_marker_cleanup(wired, monkeypatch):
    """交付文档里 ##SECTION 标记已清、锚点仍在（二者不是同一批元素）。"""
    exporter = NoteWordExporter(db=None)

    async def _load_notes(self, pid, year, sections):
        return [_fake_note(c) for c in SOE_CODES]

    monkeypatch.setattr(NoteWordExporter, "_load_notes", _load_notes)

    buf, _meta = await _run_export(exporter)
    texts = _visible_paragraphs(buf)
    assert not [t for t in texts if "##SECTION" in t]
    buf.seek(0)
    assert len(scan_anchor_blocks(Document(buf))) == len(SOE_CODES)


# ─── Property 2：可见内容不变 ────────────────────────────────────────────────


@pytest.mark.asyncio
async def test_property_2_anchor_write_does_not_change_visible_text(
    wired, monkeypatch
):
    # Feature: deliverable-lineage-wiring-and-writeback-closure, Property 2
    """开启/跳过锚点写入两次导出，可见段落文字序列逐字相等。"""
    exporter = NoteWordExporter(db=None)

    async def _load_notes(self, pid, year, sections):
        return [_fake_note(c) for c in SOE_CODES]

    monkeypatch.setattr(NoteWordExporter, "_load_notes", _load_notes)

    buf_with, _ = await _run_export(exporter)
    text_with = _visible_paragraphs(buf_with)

    # 跳过锚点写入（模拟接线前）：让 write_section_anchors 变 no-op
    monkeypatch.setattr(
        "app.services.section_anchor_utils.write_section_anchors",
        lambda doc, blocks, **kw: {},
    )
    buf_without, meta_without = await _run_export(exporter)
    text_without = _visible_paragraphs(buf_without)

    assert text_with == text_without, "锚点写入改变了可见文字"

    # 反向自检：跳过写入后文档确实没有锚点 → 证明该步骤不可省
    buf_without.seek(0)
    assert scan_anchor_blocks(Document(buf_without)) == []
    assert meta_without.anchor_map == {}


# ─── programmatic 模式零影响 ─────────────────────────────────────────────────


@pytest.mark.asyncio
async def test_programmatic_mode_meta_is_empty(monkeypatch):
    """programmatic 模式无 SECTION 块 → meta 为空壳，行为与接线前一致。"""
    exporter = NoteWordExporter(db=None)

    async def _fake_export(self, *a, **kw):
        return BytesIO(b"PK-fake")

    monkeypatch.setattr(NoteWordExporter, "export", _fake_export)
    buf, meta = await NoteWordExporter.export_with_meta(
        exporter, uuid4(), 2025, mode="programmatic"
    )
    assert buf.getvalue() == b"PK-fake"
    assert meta.kept_codes == []
    assert meta.anchor_map == {}
    assert meta.rendered_block_hashes == {}


# ─── 兼容薄壳 ────────────────────────────────────────────────────────────────


@pytest.mark.asyncio
async def test_legacy_export_template_mode_still_returns_bytesio(wired, monkeypatch):
    """`_export_template_mode` 保持只返回 BytesIO（既有调用方/测试零改动）。"""
    exporter = NoteWordExporter(db=None)

    async def _load_notes(self, pid, year, sections):
        return [_fake_note(c) for c in SOE_CODES]

    monkeypatch.setattr(NoteWordExporter, "_load_notes", _load_notes)

    out = await exporter._export_template_mode(
        uuid4(),
        2025,
        template_type="soe",
        report_scope="standalone",
        sections=None,
    )
    assert isinstance(out, BytesIO)
    out.seek(0)
    assert len(scan_anchor_blocks(Document(out))) == len(SOE_CODES)


# ─── Task 5.1：锚点 × 内容控件在生产导出路径上并存 ──────────────────────────


def _sdt_tags(doc: Document) -> list[str]:
    from docx.oxml.ns import qn

    tags: list[str] = []
    for sdt in doc.element.body.iter(qn("w:sdt")):
        pr = sdt.find(qn("w:sdtPr"))
        if pr is None:
            continue
        tag_el = pr.find(qn("w:tag"))
        if tag_el is not None:
            val = tag_el.get(qn("w:val"))
            if val:
                tags.append(val)
    return tags


@pytest.mark.asyncio
async def test_property_6_content_control_on_off_visible_text_identical(
    wired, monkeypatch
):
    # Feature: deliverable-lineage-wiring-and-writeback-closure, Property 6
    """灰度开/关两次导出：可见文字逐字相等；关闭态 sdt 数为 0；开启态锚点仍可定位。"""
    exporter = NoteWordExporter(db=None)

    async def _load_notes(self, pid, year, sections):
        return [_fake_note(c) for c in SOE_CODES]

    monkeypatch.setattr(NoteWordExporter, "_load_notes", _load_notes)

    monkeypatch.setattr(
        nwe.settings, "DELIVERABLE_LINEAGE_CONTENT_CONTROL_ENABLED", False
    )
    buf_off, _ = await _run_export(exporter)
    text_off = _visible_paragraphs(buf_off)
    buf_off.seek(0)
    assert _sdt_tags(Document(buf_off)) == [], "关闭态不得注入内容控件"

    monkeypatch.setattr(
        nwe.settings, "DELIVERABLE_LINEAGE_CONTENT_CONTROL_ENABLED", True
    )
    buf_on, meta_on = await _run_export(exporter)
    buf_on.seek(0)
    doc_on = Document(buf_on)

    # 开启态：内容控件 Tag == anchor_name，且锚点区间仍能定位到每一节
    tags = _sdt_tags(doc_on)
    for code in SOE_CODES:
        assert anchor_name(code) in tags
    blocks = scan_anchor_blocks(doc_on)
    assert {b.section_code for b in blocks} == set(SOE_CODES)
    assert set(meta_on.anchor_map) == set(SOE_CODES)

    # 可见文字（含 sdt 内部段落）与关闭态逐字相等
    buf_on.seek(0)
    text_on = _visible_paragraphs(buf_on)
    assert text_on == text_off, "内容控件注入改变了可见文字"


@pytest.mark.asyncio
async def test_content_control_on_gives_sdt_content_as_refresh_host(wired, monkeypatch):
    """开启内容控件后，刷新的插入宿主应是 w:sdtContent（新内容才被 Tag 覆盖）。"""
    from docx.oxml.ns import qn

    from app.services.section_anchor_utils import anchor_block_content_host

    exporter = NoteWordExporter(db=None)

    async def _load_notes(self, pid, year, sections):
        return [_fake_note(c) for c in SOE_CODES]

    monkeypatch.setattr(NoteWordExporter, "_load_notes", _load_notes)
    monkeypatch.setattr(
        nwe.settings, "DELIVERABLE_LINEAGE_CONTENT_CONTROL_ENABLED", True
    )

    buf, _ = await _run_export(exporter)
    buf.seek(0)
    block = scan_anchor_blocks(Document(buf))[0]
    host, children = anchor_block_content_host(block)
    assert host.tag == qn("w:sdtContent")
    assert len(children) >= 1

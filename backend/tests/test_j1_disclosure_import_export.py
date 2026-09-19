"""J1 附注披露多区块导入导出契约测试。

Spec: .kiro/specs/j1-disclosure-template-alignment/ Task 16.2

四类断言（沿用 `test_f2_disclosure_import_export.py` 范式）：
1. **前后端常量镜像**（防双真源漂移）：后端行骨架 / 列头 / 说明域 / item 键
   必须与前端 `.vue` / `.ts` 源码逐条一致
2. **workbook 结构**：模板 sheet 齐备、表头正确、行标识列在首列
3. **导出 → 导入 往返**：写进去的值能原样读回来；**并含「导出模板 → 立刻用导入
   校验器验一遍」的自检**（openpyxl 合并单元格会清空被合并格，两行表头会让
   导入自家模板报「缺少列」）
4. **派生量不采信 + 空表跳过**：期末列 / 派生父行按底稿同口径重算；
   整表未填金额时 `filled=0` 让上层跳过写库
"""
from __future__ import annotations

import io
import json
import re
from pathlib import Path

import pytest
from openpyxl import load_workbook

from app.routers.wp_render_strategies import _j1_disclosure_import_export as mod

ROOT = Path(__file__).resolve().parents[2]
FRONTEND = ROOT / "audit-platform" / "frontend" / "src"
TAB_DIR = FRONTEND / "components" / "workpaper" / "j1" / "core"
COMPOSABLES = FRONTEND / "components" / "workpaper" / "composables"
J1_COMPOSABLES = FRONTEND / "composables" / "workpaper" / "j1"

_TAB_FILE = {"listed": "J1TabDisclosureListed.vue", "soe": "J1TabDisclosureSoe.vue"}
_VARIANTS = ["listed", "soe"]


def _read(path: Path) -> str:
    return path.read_text(encoding="utf-8")


def _tab_src(variant: str) -> str:
    return _read(TAB_DIR / _TAB_FILE[variant])


def _parse_default_rows(src: str, const_name: str) -> list[tuple[str, str, int]]:
    """从 `.vue` 提取 `DEFAULT_*` 的 (id, label, indent) 序列。"""
    start = src.index(f"const {const_name}")
    end = src.index("\n]", start)
    block = src[start:end]
    out: list[tuple[str, str, int]] = []
    for m in re.finditer(r"\{\s*id:\s*'([^']+)',\s*label:\s*'([^']*)',([^}]*)\}", block):
        indent = 1 if re.search(r"indent:\s*(\d+)", m.group(3)) else 0
        out.append((m.group(1), m.group(2), indent))
    return out


def _parse_note_fields(src: str, const_name: str) -> list[tuple[str, str]]:
    """从 `j1NoteSectionMap.ts` 提取 (key, title) 序列。"""
    start = src.index(f"export const {const_name}")
    end = src.index("\n]", start)
    block = src[start:end]
    return [
        (m.group(1), m.group(2))
        for m in re.finditer(r"key:\s*'([^']+)',\s*\n?\s*title:\s*'([^']+)'", block)
    ]


# ════════════════════════════════════════════════════════════════════
# 1. 前后端常量镜像
# ════════════════════════════════════════════════════════════════════


@pytest.mark.parametrize(
    "variant,const_name,suffix",
    [
        ("listed", "DEFAULT_SUMMARY", "summary"),
        ("listed", "DEFAULT_SHORT_TERM", "short-term"),
        ("listed", "DEFAULT_POST", "post-employment"),
        ("soe", "DEFAULT_SUMMARY", "summary"),
        ("soe", "DEFAULT_SHORT_TERM", "short-term"),
        ("soe", "DEFAULT_POST", "post-employment"),
    ],
)
def test_skeleton_mirrors_frontend_defaults(variant: str, const_name: str, suffix: str):
    fe = _parse_default_rows(_tab_src(variant), const_name)
    assert fe, f"{variant} {const_name} 解析失败（.vue 格式变了？）"
    blk = next(b for b in mod.blocks_of(variant) if b.suffix == suffix)  # type: ignore[arg-type]
    assert list(blk.skeleton) == fe


@pytest.mark.parametrize("variant", _VARIANTS)
def test_category_mirrors_frontend(variant: str):
    """行 `category` 字面必须与前端一致（合计口径按 category 归组）。"""
    src = _tab_src(variant)
    for blk in mod.blocks_of(variant):  # type: ignore[arg-type]
        assert f"category: '{blk.category}'" in src


@pytest.mark.parametrize("variant", _VARIANTS)
def test_period_labels_mirror_frontend_movement_table_props(variant: str):
    """列头「期初 / 期末」用语取自底稿 UI（两变体不同），不得写死一份共用。"""
    src = _tab_src(variant)
    labels = mod._PERIOD_LABELS[variant]
    assert f'begin-label="{labels["begin"]}"' in src
    assert f'end-label="{labels["end"]}"' in src


def test_period_labels_differ_between_variants():
    """反向断言：上市/国企列头用语本就不同 → 一份常量共用即错位。"""
    assert mod._PERIOD_LABELS["listed"] != mod._PERIOD_LABELS["soe"]


@pytest.mark.parametrize(
    "variant,const_name",
    [("listed", "J1_LISTED_NOTE_FIELDS"), ("soe", "J1_SOE_NOTE_FIELDS")],
)
def test_texts_mirror_frontend_note_fields(variant: str, const_name: str):
    fe = _parse_note_fields(_read(COMPOSABLES / "j1NoteSectionMap.ts"), const_name)
    assert fe, f"{const_name} 解析失败（.ts 格式变了？）"
    assert list(mod.texts_of(variant)) == fe  # type: ignore[arg-type]


def test_item_ids_mirror_frontend_storage_keys():
    """item_id 必须与前端存储键一致，否则导入写到孤儿键。"""
    src = _read(J1_COMPOSABLES / "useJ1DisclosureSections.ts")
    assert "const prefix = `J1-disc-${variant}`" in src
    for variant in _VARIANTS:
        for blk in mod.blocks_of(variant):  # type: ignore[arg-type]
            assert f"`${{prefix}}-{blk.suffix}`" in src, f"{blk.suffix} 不在前端存储键中"
        assert f"`${{prefix}}-{mod.NOTES_SUFFIX}`" in src
        assert mod.item_id_of(variant, "summary") == f"J1-disc-{variant}-summary"  # type: ignore[arg-type]


def test_row_fields_mirror_frontend_serialize():
    """落库字段集必须与前端 `serialize()` 一致（多写 `endBalance` = 持久化派生列）。"""
    src = _read(J1_COMPOSABLES / "useJ1DisclosureSections.ts")
    start = src.index("function serialize(")
    block = src[start : src.index("function buildItems", start)]
    fe = {m.group(1) for m in re.finditer(r"^\s*(\w+):\s*r\.\w+", block, re.MULTILINE)}
    assert fe, "前端 serialize 字段解析失败"
    assert set(mod._ROW_FIELDS) == fe
    assert "endBalance" not in mod._ROW_FIELDS


def test_item_ids_unique_within_one_import():
    """同一批不得出现重复 item_id（重复会让后端整批被拒、数据全丢）。"""
    for variant in _VARIANTS:
        ids = [mod.item_id_of(variant, b.suffix) for b in mod.blocks_of(variant)]  # type: ignore[arg-type]
        ids.append(mod.item_id_of(variant, mod.NOTES_SUFFIX))  # type: ignore[arg-type]
        assert len(ids) == len(set(ids))


def test_dispatch_wired_into_existing_j1_routes():
    """不新建 router：披露 sheet_type 必须挂在既有 J1 三路由上。"""
    src = _read(ROOT / "backend/app/routers/wp_render_strategies/_j1_import_export.py")
    assert "_j1_disclosure_import_export" in src
    for st in mod.DISCLOSURE_SHEET_TYPES:
        assert st in mod.DISCLOSURE_FILE_LABELS
    assert "_DISC_SHEET_TYPES" in src
    assert src.count("in _DISC_SHEET_TYPES") == 3  # 模板 / 数据 / 导入 三处分派
    from app.routers.wp_render_strategies import _j1_import_export as route_mod

    assert set(mod.DISCLOSURE_SHEET_TYPES) <= route_mod.SHEET_TYPES


# ════════════════════════════════════════════════════════════════════
# 2. workbook 结构
# ════════════════════════════════════════════════════════════════════


@pytest.mark.parametrize("variant", _VARIANTS)
def test_template_sheets(variant: str):
    wb = mod.build_disclosure_workbook(variant)  # type: ignore[arg-type]
    names = wb.sheetnames
    assert names[0] == mod._GUIDE_SHEET
    assert mod._TEXT_SHEET in names
    assert len(mod.blocks_of(variant)) == 3  # type: ignore[arg-type]
    for blk in mod.blocks_of(variant):  # type: ignore[arg-type]
        assert blk.sheet in names
    for n in names:
        assert len(n) <= 31, f"sheet 名超 31 字符: {n}"
        assert not set(n) & set("[]:*?/\\"), f"sheet 名含非法字符: {n}"


@pytest.mark.parametrize("variant", _VARIANTS)
def test_template_headers_and_skeleton(variant: str):
    wb = mod.build_disclosure_workbook(variant)  # type: ignore[arg-type]
    headers = mod.headers_of(variant)  # type: ignore[arg-type]
    assert headers[0] == mod._KEY_HEADER, "行标识必须在首列（行名可改，不能按名匹配）"
    for blk in mod.blocks_of(variant):  # type: ignore[arg-type]
        ws = wb[blk.sheet]
        hdr = mod._find_header_row(ws, headers[0])
        assert hdr is not None, f"{blk.sheet} 找不到表头行"
        actual = [
            str(ws.cell(row=hdr, column=c).value or "") for c in range(1, len(headers) + 1)
        ]
        assert actual == list(headers)
        keys = [ws.cell(row=hdr + 1 + i, column=1).value for i in range(len(blk.skeleton))]
        labels = [ws.cell(row=hdr + 1 + i, column=2).value for i in range(len(blk.skeleton))]
        indents = [ws.cell(row=hdr + 1 + i, column=3).value for i in range(len(blk.skeleton))]
        assert keys == [rk for rk, _lb, _i in blk.skeleton]
        assert labels == [lb for _rk, lb, _i in blk.skeleton]
        assert indents == [i for _rk, _lb, i in blk.skeleton]


@pytest.mark.parametrize("variant", _VARIANTS)
def test_end_column_marked_derived_and_not_importable(variant: str):
    headers = mod.headers_of(variant)  # type: ignore[arg-type]
    assert headers[-1].endswith(mod._DERIVED_SUFFIX)
    assert mod._PERIOD_LABELS[variant]["end"] in headers[-1]
    # 期末列的**纯列名**不在可导入的数值列里
    assert mod._PERIOD_LABELS[variant]["end"] not in mod._numeric_headers(variant)  # type: ignore[arg-type]


@pytest.mark.parametrize("variant", _VARIANTS)
def test_text_sheet_lists_all_note_blocks(variant: str):
    wb = mod.build_disclosure_workbook(variant)  # type: ignore[arg-type]
    ws = wb[mod._TEXT_SHEET]
    keys = [ws.cell(row=r, column=1).value for r in range(3, ws.max_row + 1)]
    titles = [ws.cell(row=r, column=2).value for r in range(3, ws.max_row + 1)]
    assert keys == [k for k, _t in mod.texts_of(variant)]  # type: ignore[arg-type]
    assert titles == [t for _k, t in mod.texts_of(variant)]  # type: ignore[arg-type]


@pytest.mark.parametrize("variant", _VARIANTS)
def test_guidance_documents_hard_constraints(variant: str):
    lines = "\n".join(mod._guidance_lines(variant))  # type: ignore[arg-type]
    assert mod._KEY_HEADER in lines
    assert "整表覆盖" in lines
    assert "公式" in lines
    for blk in mod.blocks_of(variant):  # type: ignore[arg-type]
        assert blk.sheet in lines
        assert "整表覆盖" in blk.note


@pytest.mark.parametrize("variant", _VARIANTS)
def test_soe_and_listed_sheet_names_reflect_source_sections(variant: str):
    """国企侧三张表都是「…列示」（源模板 / 附注表名口径），上市侧不是。"""
    sheets = [b.sheet for b in mod.blocks_of(variant)]  # type: ignore[arg-type]
    if variant == "soe":
        assert all(s.endswith("列示") for s in sheets)
    else:
        assert not any(s.endswith("列示") for s in sheets)


# ════════════════════════════════════════════════════════════════════
# 3. 往返（纯解析层，不碰 DB）
# ════════════════════════════════════════════════════════════════════


def _wb_bytes(wb) -> bytes:
    buf = io.BytesIO()
    wb.save(buf)
    return buf.getvalue()


def _reload(wb):
    return load_workbook(io.BytesIO(_wb_bytes(wb)), data_only=True)


def _set(ws, row_key: str, col_name: str, value):
    hdr = mod._find_header_row(ws, mod._KEY_HEADER)
    idx = mod._header_index(ws, hdr, 7)
    for r in range(hdr + 1, ws.max_row + 1):
        if str(ws.cell(row=r, column=1).value or "").strip() == row_key:
            ws.cell(row=r, column=idx[col_name]).value = value
            return r
    raise AssertionError(f"未找到行 {row_key}")


@pytest.mark.parametrize("variant", _VARIANTS)
def test_template_roundtrips_through_importer(variant: str):
    """自检：导出模板 → 立刻用导入校验器验一遍（表头必须能被找到、列不缺）。

    这条守的是 openpyxl 合并单元格坑：若把标签列做成两行纵向合并表头，
    第 2 行会被清空 → 导入自家模板必报「缺少列」。
    """
    wb = _reload(mod.build_disclosure_workbook(variant))  # type: ignore[arg-type]
    for blk in mod.blocks_of(variant):  # type: ignore[arg-type]
        rows, warns, filled = mod._import_rows(wb[blk.sheet], blk, variant)  # type: ignore[arg-type]
        assert not any("未找到表头" in w for w in warns), warns
        assert not any("缺少列" in w for w in warns), warns
        # 空模板：骨架行被读到但一个金额都没填 → 不得写库
        assert filled == 0
        assert [r["id"] for r in rows] == [rk for rk, _lb, _i in blk.skeleton]
    texts, warns, cnt = mod._import_texts(wb[mod._TEXT_SHEET], variant)  # type: ignore[arg-type]
    assert not any("未找到表头" in w for w in warns), warns
    assert cnt == 0 and texts == {}  # 空模板不清既有说明


@pytest.mark.parametrize("variant", _VARIANTS)
def test_roundtrip_rows_values(variant: str):
    wb = mod.build_disclosure_workbook(variant)  # type: ignore[arg-type]
    blk = next(b for b in mod.blocks_of(variant) if b.suffix == "summary")  # type: ignore[arg-type]
    begin_h, inc_h, dec_h = mod._numeric_headers(variant)  # type: ignore[arg-type]
    ws = wb[blk.sheet]
    _set(ws, "s-1", begin_h, 1000)
    _set(ws, "s-1", inc_h, 250.5)
    _set(ws, "s-1", dec_h, 100)
    _set(ws, "s-2", begin_h, 30)

    rows, warns, filled = mod._import_rows(_reload(wb)[blk.sheet], blk, variant)  # type: ignore[arg-type]
    assert warns == []
    assert filled == 2
    by_id = {r["id"]: r for r in rows}
    assert by_id["s-1"]["beginBalance"] == 1000.0
    assert by_id["s-1"]["increase"] == 250.5
    assert by_id["s-1"]["endBalance"] == 1150.5  # 期末派生
    assert by_id["s-2"]["endBalance"] == 30.0
    assert by_id["s-1"]["category"] == blk.category


@pytest.mark.parametrize("variant", _VARIANTS)
def test_import_recalculates_end_balance_ignoring_typed_value(variant: str):
    """期末列是派生列：用户填的值不作权威（防导入引入不一致数据）。"""
    wb = mod.build_disclosure_workbook(variant)  # type: ignore[arg-type]
    blk = next(b for b in mod.blocks_of(variant) if b.suffix == "summary")  # type: ignore[arg-type]
    begin_h, _inc, _dec = mod._numeric_headers(variant)  # type: ignore[arg-type]
    headers = mod.headers_of(variant)  # type: ignore[arg-type]
    ws = wb[blk.sheet]
    _set(ws, "s-1", begin_h, 100)
    _set(ws, "s-1", headers[-1], 999999)  # 乱填期末

    rows, _w, _f = mod._import_rows(_reload(wb)[blk.sheet], blk, variant)  # type: ignore[arg-type]
    assert next(r for r in rows if r["id"] == "s-1")["endBalance"] == 100.0


@pytest.mark.parametrize("variant", _VARIANTS)
def test_import_recomputes_parent_rows_from_indented_children(variant: str):
    """派生父行 = Σ 紧邻缩进子行（源公式 SUM）；用户在父行乱填不作准。"""
    wb = mod.build_disclosure_workbook(variant)  # type: ignore[arg-type]
    blk = next(b for b in mod.blocks_of(variant) if b.suffix == "short-term")  # type: ignore[arg-type]
    begin_h, inc_h, dec_h = mod._numeric_headers(variant)  # type: ignore[arg-type]
    ws = wb[blk.sheet]
    _set(ws, "st-3", begin_h, 88888)  # 社会保险费（父行，公式）
    _set(ws, "st-4", begin_h, 100)
    _set(ws, "st-4", inc_h, 10)
    _set(ws, "st-4", dec_h, 5)
    _set(ws, "st-5", begin_h, 50)
    _set(ws, "st-5", inc_h, 5)
    _set(ws, "st-5", dec_h, 2)

    rows, _w, _f = mod._import_rows(_reload(wb)[blk.sheet], blk, variant)  # type: ignore[arg-type]
    parent = next(r for r in rows if r["id"] == "st-3")
    assert parent["beginBalance"] == 150.0
    assert parent["increase"] == 15.0
    assert parent["decrease"] == 7.0
    assert parent["endBalance"] == 158.0


@pytest.mark.parametrize("variant", _VARIANTS)
def test_indent_survives_roundtrip(variant: str):
    """「其中：」缩进是结构信息（合计排除它）→ 往返不得丢。"""
    wb = mod.build_disclosure_workbook(variant)  # type: ignore[arg-type]
    blk = next(b for b in mod.blocks_of(variant) if b.suffix == "post-employment")  # type: ignore[arg-type]
    begin_h, _i, _d = mod._numeric_headers(variant)  # type: ignore[arg-type]
    ws = wb[blk.sheet]
    _set(ws, "pe-2", begin_h, 10)
    rows, _w, _f = mod._import_rows(_reload(wb)[blk.sheet], blk, variant)  # type: ignore[arg-type]
    got = {r["id"]: r["indent"] for r in rows}
    assert got == {rk: ind for rk, _lb, ind in blk.skeleton}


@pytest.mark.parametrize("variant", _VARIANTS)
def test_renamed_label_still_matched_by_row_key(variant: str):
    """行名可改 → 必须按行标识匹配（改名后仍落到同一 rowKey）。"""
    wb = mod.build_disclosure_workbook(variant)  # type: ignore[arg-type]
    blk = next(b for b in mod.blocks_of(variant) if b.suffix == "summary")  # type: ignore[arg-type]
    begin_h, _i, _d = mod._numeric_headers(variant)  # type: ignore[arg-type]
    ws = wb[blk.sheet]
    headers = mod.headers_of(variant)  # type: ignore[arg-type]
    _set(ws, "s-1", headers[1], "改过名的短期薪酬")
    _set(ws, "s-1", begin_h, 7)
    rows, _w, _f = mod._import_rows(_reload(wb)[blk.sheet], blk, variant)  # type: ignore[arg-type]
    row = next(r for r in rows if r["id"] == "s-1")
    assert row["label"] == "改过名的短期薪酬"
    assert row["beginBalance"] == 7.0


def test_duplicate_labels_exist_hence_row_key_required():
    """守卫：骨架里确实存在重复行名 → 「行标识」列不可省。"""
    labels = [lb for _rk, lb, _i in mod._SOE_SHORT_TERM_SKELETON] + [
        lb for _rk, lb, _i in mod._SOE_POST_SKELETON
    ]
    assert len([lb for lb in labels if labels.count(lb) > 1]) >= 2


@pytest.mark.parametrize("variant", _VARIANTS)
def test_duplicate_row_key_is_skipped_with_warning(variant: str):
    wb = mod.build_disclosure_workbook(variant)  # type: ignore[arg-type]
    blk = next(b for b in mod.blocks_of(variant) if b.suffix == "summary")  # type: ignore[arg-type]
    begin_h, _i, _d = mod._numeric_headers(variant)  # type: ignore[arg-type]
    ws = wb[blk.sheet]
    _set(ws, "s-1", begin_h, 5)
    ws.cell(row=ws.max_row + 1, column=1).value = "s-1"
    ws.cell(row=ws.max_row, column=2).value = "重复行"
    rows, warns, _f = mod._import_rows(_reload(wb)[blk.sheet], blk, variant)  # type: ignore[arg-type]
    assert len([r for r in rows if r["id"] == "s-1"]) == 1
    assert any("重复" in w for w in warns)


@pytest.mark.parametrize("variant", _VARIANTS)
def test_blank_table_returns_zero_so_caller_skips(variant: str):
    """整表未填金额 → filled=0（上层跳过写库，不清既有数据）。"""
    wb = mod.build_disclosure_workbook(variant)  # type: ignore[arg-type]
    for blk in mod.blocks_of(variant):  # type: ignore[arg-type]
        _rows, warns, filled = mod._import_rows(_reload(wb)[blk.sheet], blk, variant)  # type: ignore[arg-type]
        assert filled == 0
        assert any("未改动既有数据" in w for w in warns)


@pytest.mark.parametrize("variant", _VARIANTS)
def test_new_rows_without_row_key_get_generated_id(variant: str):
    wb = mod.build_disclosure_workbook(variant)  # type: ignore[arg-type]
    blk = next(b for b in mod.blocks_of(variant) if b.suffix == "summary")  # type: ignore[arg-type]
    ws = wb[blk.sheet]
    idx_begin = mod._header_index(ws, mod._find_header_row(ws, mod._KEY_HEADER), 7)[
        mod._numeric_headers(variant)[0]  # type: ignore[arg-type]
    ]
    r = ws.max_row + 1
    ws.cell(row=r, column=2).value = "手工新增行"
    ws.cell(row=r, column=idx_begin).value = 42
    rows, _w, filled = mod._import_rows(_reload(wb)[blk.sheet], blk, variant)  # type: ignore[arg-type]
    assert filled == 1
    new = next(r_ for r_ in rows if r_["label"] == "手工新增行")
    assert new["id"].startswith(f"{blk.suffix}-")


@pytest.mark.parametrize("variant", _VARIANTS)
def test_roundtrip_texts(variant: str):
    wb = mod.build_disclosure_workbook(variant)  # type: ignore[arg-type]
    ws = wb[mod._TEXT_SHEET]
    first_key, _title = mod.texts_of(variant)[0]  # type: ignore[arg-type]
    for r in range(3, ws.max_row + 1):
        if str(ws.cell(row=r, column=1).value or "").strip() == first_key:
            ws.cell(row=r, column=3).value = "多行\n说明内容"
    texts, warns, cnt = mod._import_texts(_reload(wb)[mod._TEXT_SHEET], variant)  # type: ignore[arg-type]
    assert warns == []
    assert cnt == len(mod.texts_of(variant))  # type: ignore[arg-type]
    assert texts[first_key] == "多行\n说明内容"


@pytest.mark.parametrize("variant", _VARIANTS)
def test_texts_unknown_key_warns_and_title_fallback_works(variant: str):
    wb = mod.build_disclosure_workbook(variant)  # type: ignore[arg-type]
    ws = wb[mod._TEXT_SHEET]
    key, title = mod.texts_of(variant)[-1]  # type: ignore[arg-type]
    r = ws.max_row + 1
    ws.cell(row=r, column=1).value = "made-up-key"
    ws.cell(row=r, column=2).value = "不存在的区块"
    ws.cell(row=r, column=3).value = "x"
    # 用户把标识列清空但保留标题 → 回退按标题匹配
    for rr in range(3, ws.max_row + 1):
        if str(ws.cell(row=rr, column=1).value or "").strip() == key:
            ws.cell(row=rr, column=1).value = None
            ws.cell(row=rr, column=3).value = "按标题回退命中"
    texts, warns, _c = mod._import_texts(_reload(wb)[mod._TEXT_SHEET], variant)  # type: ignore[arg-type]
    assert texts.get(key) == "按标题回退命中"
    assert any("made-up-key" in w for w in warns)
    assert title  # 标题非空（回退依据）


def test_export_with_data_renders_values():
    remarks = {
        "J1-disc-listed-summary": json.dumps(
            [
                {
                    "id": "s-1",
                    "label": "短期薪酬",
                    "category": "summary",
                    "indent": 0,
                    "beginBalance": 100,
                    "increase": 20,
                    "decrease": 5,
                }
            ]
        ),
        "J1-disc-listed-notes": json.dumps({"shortTerm": "非货币性福利以食堂餐补形式发放。"}),
    }
    wb = _reload(mod.build_disclosure_workbook("listed", remarks=remarks))
    blk = next(b for b in mod._LISTED_BLOCKS if b.suffix == "summary")
    rows, _w, filled = mod._import_rows(wb[blk.sheet], blk, "listed")
    assert filled == 1
    assert rows[0]["beginBalance"] == 100.0
    assert rows[0]["endBalance"] == 115.0
    texts, _w2, _c = mod._import_texts(wb[mod._TEXT_SHEET], "listed")
    assert texts["shortTerm"] == "非货币性福利以食堂餐补形式发放。"


def test_malformed_remark_json_falls_back_to_skeleton():
    remarks = {"J1-disc-soe-summary": "{ not json", "J1-disc-soe-notes": "[]"}
    wb = mod.build_disclosure_workbook("soe", remarks=remarks)
    ws = wb[mod._SOE_BLOCKS[0].sheet]
    hdr = mod._find_header_row(ws, mod._KEY_HEADER)
    keys = [ws.cell(row=hdr + 1 + i, column=1).value for i in range(len(mod._SOE_SUMMARY_SKELETON))]
    assert keys == [rk for rk, _lb, _i in mod._SOE_SUMMARY_SKELETON]
    assert mod._TEXT_SHEET in wb.sheetnames


def test_header_row_tolerates_reordered_columns():
    from openpyxl import Workbook as _Wb

    blk = next(b for b in mod._LISTED_BLOCKS if b.suffix == "summary")
    wb = _Wb()
    ws = wb.active
    ws.title = blk.sheet
    ws.append(["提示：xxx"])
    ws.append([mod._KEY_HEADER, "项目", "本期减少", "上年年末数", "本期增加"])
    ws.append(["s-1", "短期薪酬", 5, 100, 20])
    rows, _w, filled = mod._import_rows(_reload(wb)[blk.sheet], blk, "listed")
    assert filled == 1
    assert rows[0]["beginBalance"] == 100.0
    assert rows[0]["decrease"] == 5.0
    assert rows[0]["endBalance"] == 115.0


def test_two_row_merged_header_still_readable():
    """回归守卫：纵向合并会清空第 2 行（openpyxl 语义）→ 读表头须回退上一行。"""
    from openpyxl import Workbook as _Wb

    blk = next(b for b in mod._LISTED_BLOCKS if b.suffix == "summary")
    wb = _Wb()
    ws = wb.active
    ws.title = blk.sheet
    ws.append([mod._KEY_HEADER, "项目", "本期发生额", None])
    ws.append([None, None, "上年年末数", "本期增加"])
    ws.merge_cells("A1:A2")
    ws.merge_cells("B1:B2")
    ws.append(["s-1", "短期薪酬", 100, 20])
    ws2 = _reload(wb)[blk.sheet]
    idx = mod._header_index(ws2, 2, 4)
    assert idx[mod._KEY_HEADER] == 1  # 被合并清空的 A2 回退取 A1
    assert idx["项目"] == 2
    assert idx["上年年末数"] == 3


def test_missing_sheet_is_skipped_not_error():
    from openpyxl import Workbook as _Wb

    wb = _Wb()
    wb.active.title = mod._TEXT_SHEET
    ws = wb.active
    ws.append(["提示"])
    ws.append(list(mod._TEXT_HEADERS))
    key, title = mod.texts_of("listed")[0]
    ws.append([key, title, "仅导入文本"])
    texts, _w, cnt = mod._import_texts(_reload(wb)[mod._TEXT_SHEET], "listed")
    assert cnt == 1
    assert texts[key] == "仅导入文本"


def test_row_limit_truncates_with_warning():
    from openpyxl import Workbook as _Wb

    blk = next(b for b in mod._LISTED_BLOCKS if b.suffix == "summary")
    wb = _Wb()
    ws = wb.active
    ws.title = blk.sheet
    ws.append(list(mod.headers_of("listed")))
    for i in range(mod.ROW_LIMIT + 10):
        ws.append([f"x-{i}", f"行{i}", 0, 1, 0, 0, None])
    rows, warns, _f = mod._import_rows(_reload(wb)[blk.sheet], blk, "listed")
    assert len(rows) == mod.ROW_LIMIT
    assert any("截断" in w for w in warns)


def test_serialize_rows_matches_frontend_shape():
    rows = mod.apply_derivations(
        [
            {
                "id": "s-1",
                "label": "短期薪酬",
                "category": "summary",
                "indent": 0,
                "beginBalance": 10,
                "increase": 5,
                "decrease": 2,
            }
        ]
    )
    payload = json.loads(mod.serialize_rows(rows))
    assert list(payload[0].keys()) == list(mod._ROW_FIELDS)
    assert "endBalance" not in payload[0]


# ════════════════════════════════════════════════════════════════════
# 4. 路由分派入参归一（共享前端组件 POST + `sheet`；历史 composable GET + `sheet_type`）
# ════════════════════════════════════════════════════════════════════


def test_routes_accept_sheet_and_sheet_type_over_get_and_post():
    """`CycleImportExportDropdown` 走 POST + `sheet`；J1 历史 composable 走 GET + `sheet_type`。"""
    import inspect

    from app.routers.wp_render_strategies import _j1_import_export as route_mod

    methods: dict[str, set[str]] = {}
    for r in route_mod.router.routes:
        methods.setdefault(r.path, set()).update(getattr(r, "methods", set()) or set())
    for name in ("export-template", "export-data", "import-data"):
        path = f"/api/workpapers/{{wp_id}}/j1/{name}"
        assert path in methods, f"缺路由 {path}"
        assert "POST" in methods[path], f"{path} 必须支持 POST（共享前端组件只发 POST）"
    for name in ("export-template", "export-data"):
        assert "GET" in methods[f"/api/workpapers/{{wp_id}}/j1/{name}"], "历史 GET 调用不得失效"

    for fn in (route_mod.export_template, route_mod.export_data, route_mod.import_data):
        params = inspect.signature(fn).parameters
        assert "sheet" in params and "sheet_type" in params


def test_resolve_sheet_type_normalizes_and_rejects():
    from fastapi import HTTPException

    from app.routers.wp_render_strategies._j1_import_export import _resolve_sheet_type

    assert _resolve_sheet_type(mod.SHEET_LISTED, None) == mod.SHEET_LISTED
    assert _resolve_sheet_type(None, "detail") == "detail"
    # 两者同时给出时 `sheet` 优先（共享组件传的就是它）
    assert _resolve_sheet_type(mod.SHEET_SOE, "detail") == mod.SHEET_SOE
    for bad_sheet, bad_type in ((None, None), ("", ""), ("made-up", None)):
        with pytest.raises(HTTPException):
            _resolve_sheet_type(bad_sheet, bad_type)


# ════════════════════════════════════════════════════════════════════
# 5. 写库路径（伪 session，验证 item_id / 载荷 / 全空跳过）
# ════════════════════════════════════════════════════════════════════


class _FakeResult:
    def __init__(self, scalar=None, rows=None):
        self._scalar = scalar
        self._rows = rows or []

    def scalar_one_or_none(self):
        return self._scalar

    def all(self):
        return self._rows


class _FakeDB:
    """只实现 `import_disclosure_data` 用到的 execute / commit。"""

    def __init__(self, remarks: dict[str, str] | None = None):
        self.remarks = dict(remarks or {})
        self.writes: list[tuple[str, str]] = []
        self.commits = 0

    async def execute(self, stmt, params=None):  # noqa: ANN001
        sql = " ".join(str(stmt).split())
        params = params or {}
        if "FROM working_paper" in sql:
            return _FakeResult(scalar="proj-1")
        if sql.startswith("SELECT item_id, remark"):
            pat = str(params.get("pat", "")).rstrip("%")
            return _FakeResult(
                rows=[(k, v) for k, v in self.remarks.items() if k.startswith(pat)]
            )
        if "INSERT INTO checklist_responses" in sql:
            self.writes.append((params["item_id"], params["remark"]))
            self.remarks[params["item_id"]] = params["remark"]
            return _FakeResult()
        raise AssertionError(f"未预期的 SQL: {sql[:120]}")

    async def commit(self):
        self.commits += 1


def _fill_all_blocks(wb, variant: str) -> None:
    begin_h, inc_h, dec_h = mod._numeric_headers(variant)  # type: ignore[arg-type]
    for blk in mod.blocks_of(variant):  # type: ignore[arg-type]
        first_key = blk.skeleton[0][0]
        _set(wb[blk.sheet], first_key, begin_h, 100)
        _set(wb[blk.sheet], first_key, inc_h, 20)
        _set(wb[blk.sheet], first_key, dec_h, 5)


def _fill_first_text(wb, variant: str, content: str) -> None:
    ws = wb[mod._TEXT_SHEET]
    key = mod.texts_of(variant)[0][0]  # type: ignore[arg-type]
    for r in range(3, ws.max_row + 1):
        if str(ws.cell(row=r, column=1).value or "").strip() == key:
            ws.cell(row=r, column=3).value = content


@pytest.mark.parametrize("variant", _VARIANTS)
def test_import_writes_expected_item_ids_and_payload(variant: str):
    import asyncio

    wb = mod.build_disclosure_workbook(variant)  # type: ignore[arg-type]
    _fill_all_blocks(wb, variant)
    _fill_first_text(wb, variant, "说明正文")
    db = _FakeDB()
    result = asyncio.run(mod.import_disclosure_data(db, "wp-1", variant, _wb_bytes(wb)))  # type: ignore[arg-type]

    assert result["ok"] is True
    assert db.commits == 1
    written = dict(db.writes)
    for blk in mod.blocks_of(variant):  # type: ignore[arg-type]
        iid = mod.item_id_of(variant, blk.suffix)  # type: ignore[arg-type]
        assert iid in written
        payload = json.loads(written[iid])
        assert list(payload[0].keys()) == list(mod._ROW_FIELDS)
        assert payload[0]["category"] == blk.category
        assert len(payload) == len(blk.skeleton)
    notes = json.loads(written[mod.item_id_of(variant, mod.NOTES_SUFFIX)])  # type: ignore[arg-type]
    assert notes[mod.texts_of(variant)[0][0]] == "说明正文"  # type: ignore[arg-type]


@pytest.mark.parametrize("variant", _VARIANTS)
def test_import_blank_template_writes_nothing(variant: str):
    """拿空模板导入 → 一个字节都不写库（防误清既有明细与说明）。"""
    import asyncio

    wb = mod.build_disclosure_workbook(variant)  # type: ignore[arg-type]
    db = _FakeDB({mod.item_id_of(variant, "summary"): "[{\"id\":\"s-1\"}]"})  # type: ignore[arg-type]
    result = asyncio.run(mod.import_disclosure_data(db, "wp-1", variant, _wb_bytes(wb)))  # type: ignore[arg-type]

    assert result["ok"] is False
    assert db.writes == []
    assert db.commits == 0


@pytest.mark.parametrize("variant", _VARIANTS)
def test_import_texts_only_does_not_touch_row_blocks(variant: str):
    import asyncio

    wb = mod.build_disclosure_workbook(variant)  # type: ignore[arg-type]
    _fill_first_text(wb, variant, "只导文本")
    db = _FakeDB()
    result = asyncio.run(mod.import_disclosure_data(db, "wp-1", variant, _wb_bytes(wb)))  # type: ignore[arg-type]

    assert result["ok"] is True
    assert [iid for iid, _ in db.writes] == [mod.item_id_of(variant, mod.NOTES_SUFFIX)]  # type: ignore[arg-type]


def test_import_notes_merge_preserves_keys_absent_from_sheet():
    """「文本说明」只留一行时，其余说明域沿用既有值（不被清空）。"""
    import asyncio

    from openpyxl import Workbook as _Wb

    variant = "listed"
    key, title = mod.texts_of(variant)[0]
    other_key = mod.texts_of(variant)[1][0]
    wb = _Wb()
    ws = wb.active
    ws.title = mod._TEXT_SHEET
    ws.append(["提示"])
    ws.append(list(mod._TEXT_HEADERS))
    ws.append([key, title, "新的短期薪酬说明"])
    db = _FakeDB(
        {mod.item_id_of(variant, mod.NOTES_SUFFIX): json.dumps({other_key: "原有设定提存说明"})}
    )
    asyncio.run(mod.import_disclosure_data(db, "wp-1", variant, _wb_bytes(wb)))

    notes = json.loads(dict(db.writes)[mod.item_id_of(variant, mod.NOTES_SUFFIX)])
    assert notes[key] == "新的短期薪酬说明"
    assert notes[other_key] == "原有设定提存说明"

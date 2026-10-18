"""附注模板模式导出的收尾清理与截断编号匹配（第 26 组）。

现象（真实项目只读导出，2026-10-16）：交付中心附注默认走模板模式，且默认把每节内容包进内容控件 ``w:sdt``
（``USE_TEMPLATE_FILL_SERVICE`` / ``DELIVERABLE_LINEAGE_CONTENT_CONTROL_ENABLED`` 均默认开）。收尾两道清理
（``remove_section_markers`` / ``_strip_residual_note_placeholders``）按 ``doc.paragraphs`` / ``doc.tables`` 遍历，
只看 body 直接子元素 ⇒ 内容控件里的 ``##STYLE_REF:…##`` 与未填充的 ``{{section:…}}`` 原样进了交付件：
甲 117 + 19、安徽 201 + 19、和平 201 + 6、首汽 218 + 0 处。``test_note_template_mode`` 两例同因（按
``doc.paragraphs`` 取文，看到的是 0 段）。
另：模板块编号是旧种子的 10 字截断码（「四、记账基础和计价原」），新生成项目是全码 ⇒ 甲 / 乙 / 安徽各 12 章匹配
不到附注，交付件里只有模板标题与占位符。索引里有 section_id，与附注 section_id 一致（8 个项目逐块对照 0 冲突）。

判据：
- 内容控件开 / 关两种情况下，导出件全部段落（含内容控件与表格单元格）都没有 ``##SECTION`` / ``##STYLE_REF`` /
  ``{{section|table|seq:…}}``；开时导出件里确有 ``w:sdt``（否则判据空转）；
- 单元格里唯一的段落只清空不删（``w:tc`` 至少要一个段落，否则 Word 报内容损坏）；
- 块编号与附注编号不同、section_id 相同 ⇒ 按 section_id 填入；section_id 在附注里不唯一、或该附注已被别的块按
  编号认领 ⇒ 不猜；
- 真实模板（soe_standalone）：给截断编号章节一条全码附注，导出件里有其正文且无任何残留标记。
"""
from __future__ import annotations

import re
import zipfile
from io import BytesIO
from pathlib import Path

import pytest
from docx import Document
from lxml import etree

import app.services.note_word_exporter as nwe_mod
from app.services.note_word_exporter import NoteWordExporter
from tests.services.test_note_template_mode import _make_db, _make_note, _patch_loader, _patch_scope

W = "{http://schemas.openxmlformats.org/wordprocessingml/2006/main}"
_LEFTOVER = re.compile(r"##/?(?:SECTION|STYLE_REF)[^#]*##|\{\{(?:section|table|seq):[^}]*\}\}")
TRUNC, FULL, SID = "四、记账基础和计价原", "四、记账基础和计价原则", "sid-basis"
_REAL_TEMPLATE = (Path(__file__).resolve().parents[1] / "data" / "audit_report_templates" / "disclosure_notes"
                  / "soe_standalone.docx")


def _root(data: bytes):
    with zipfile.ZipFile(BytesIO(data)) as z:
        return etree.fromstring(z.read("word/document.xml"))


def _texts(data: bytes) -> list[str]:
    """全部段落文字（文档顺序，含内容控件、表格单元格里的段落）—— python-docx 的 doc.paragraphs 看不到内容控件里的段落。"""
    return ["".join(t.text or "" for t in p.iter(W + "t")) for p in _root(data).iter(W + "p")]


def _build(tmp_path: Path) -> Path:
    doc = Document()

    def block(code: str, *body: str, cell: str | None = None) -> None:
        doc.add_paragraph(f"##SECTION:{code}##")
        for line in body:
            doc.add_paragraph(line)
        if cell is not None:
            doc.add_table(rows=1, cols=1).cell(0, 0).paragraphs[0].add_run(cell)
        doc.add_paragraph(f"##/SECTION:{code}##")

    block("八、1", "货币资金", "##STYLE_REF:table:八、1:0##", "{{section:八、1}}", cell="##STYLE_REF:table:八、1:1##")
    block("八、2", "交易性金融资产", "{{section:八、2}}", cell="{{table:八、2:0}}")  # 无附注：占位符须清掉
    block(TRUNC, "记账基础和计价原则", "{{section:%s}}" % TRUNC)
    out = tmp_path / "soe_standalone.docx"
    doc.save(str(out))
    return out


def _patch_index(monkeypatch, extra_sid: dict[str, str] | None = None) -> None:
    sid = {"八、1": "sid-cash", "八、2": "sid-tfa", TRUNC: SID, **(extra_sid or {})}
    monkeypatch.setattr(nwe_mod, "_load_section_code_index", lambda _vk: [
        {"section_code": code, "section_id": value, "legacy_aliases": []} for code, value in sid.items()])


def _note(code: str, text: str, section_id: str | None = None):
    note = _make_note(code, code, text_content=text)
    note.section_id = section_id
    return note


async def _export(notes: list) -> bytes:
    bio = await NoteWordExporter(_make_db(notes)).export(
        project_id=nwe_mod.UUID(int=1), year=2025, template_type="soe", report_scope="standalone", mode="template")
    return bio.getvalue()


@pytest.fixture
def synthetic(tmp_path, monkeypatch):
    _patch_loader(monkeypatch, _build(tmp_path))
    _patch_index(monkeypatch)
    _patch_scope(monkeypatch)
    return monkeypatch


@pytest.mark.parametrize("controls", [True, False], ids=["内容控件开", "内容控件关"])
@pytest.mark.asyncio
async def test_no_leftover_markers_or_placeholders(synthetic, controls):
    synthetic.setattr(nwe_mod.settings, "DELIVERABLE_LINEAGE_CONTENT_CONTROL_ENABLED", controls)
    data = await _export([_note("八、1", "货币资金期末余额披露", "sid-cash")])
    texts = _texts(data)
    assert "货币资金期末余额披露" in texts
    assert [t for t in texts if _LEFTOVER.search(t)] == []
    assert (len(_root(data).findall(f".//{W}sdt")) > 0) is controls, "内容控件开关没有生效，判据空转"


@pytest.mark.asyncio
async def test_table_cell_keeps_a_paragraph_after_cleanup(synthetic):
    # spec deliverable-trio-generation-accuracy：模板样例表不再原样进交付件（保留块按模板结构重建，样例表换成库内表、
    # 无库内表即删）⇒ 夹具里带标记的样例表被删，判据改为「库内表渲染后单元格都有段落」；标记段落在单元格里只清空
    # 不删的口径由 tests/services/test_note_template_block_fill.py 直接对 remove_section_markers 守卫
    note = _note("八、1", "货币资金期末余额披露", "sid-cash")
    note.table_data = {"headers": ["项目", "期末余额"], "rows": [{"label": "银行存款", "values": [100.0]}]}
    data = await _export([note])
    cells = list(_root(data).iter(W + "tc"))
    assert cells, "夹具没有表格，判据空转"
    assert all(cell.find(W + "p") is not None for cell in cells)


@pytest.mark.asyncio
async def test_truncated_block_code_is_filled_by_section_id(synthetic):
    data = await _export([_note(FULL, "本公司以权责发生制为记账基础。", SID)])
    assert "本公司以权责发生制为记账基础。" in _texts(data)


@pytest.mark.asyncio
async def test_ambiguous_section_id_is_not_guessed(synthetic):
    data = await _export([_note(FULL, "第一条。", SID), _note("四、记账基础", "第二条。", SID)])
    texts = _texts(data)
    assert "第一条。" not in texts and "第二条。" not in texts


@pytest.mark.asyncio
async def test_note_claimed_by_code_is_not_reused_by_section_id(tmp_path, monkeypatch):
    """附注「八、1」按编号归「八、1」块；另一块的索引 section_id 恰好也指向它时不得重复填入。"""
    _patch_loader(monkeypatch, _build(tmp_path))
    _patch_index(monkeypatch, {TRUNC: "sid-cash"})
    _patch_scope(monkeypatch)
    data = await _export([_note("八、1", "货币资金期末余额披露", "sid-cash")])
    assert _texts(data).count("货币资金期末余额披露") == 1


@pytest.mark.parametrize("controls", [True, False], ids=["内容控件开", "内容控件关"])
@pytest.mark.asyncio
async def test_rendered_block_hash_matches_delivered_doc(synthetic, controls):
    """导出时记的 Rendered_Block_Hash = 交付件里按锚点重算的块哈希（第 27 组）。

    收尾清理在算哈希之后：清掉的占位符若算进了哈希，没动过的块会被增量刷新的人工编辑检测、回填的
    「交付后改过」判据当成人工改过。第 26 组让清理真正生效后，真实项目每份 5 块不一致。
    """
    from app.services.section_anchor_utils import block_text_hash, resolve_section_blocks

    synthetic.setattr(nwe_mod.settings, "DELIVERABLE_LINEAGE_CONTENT_CONTROL_ENABLED", controls)
    # spec deliverable-trio-generation-accuracy：无附注的块不再按模板原貌保留（整块删除）⇒ 给「八、2」一条附注，
    # 它块内单元格里的 {{table:…}} 样例表随重建删除，判据仍是「导出时记的哈希 = 交付件按锚点重算的哈希」
    exporter = NoteWordExporter(_make_db([_note("八、1", "货币资金期末余额披露", "sid-cash"),
                                          _note("八、2", "交易性金融资产期末余额披露", "sid-tfa")]))
    buf, meta = await exporter.export_with_meta(
        nwe_mod.UUID(int=1), 2025, template_type="soe", report_scope="standalone", mode="template")
    mode, blocks = resolve_section_blocks(Document(BytesIO(buf.getvalue())))
    delivered = {b.section_code: block_text_hash(list(b.elements)) for b in blocks}
    assert mode == "anchor" and "八、2" in delivered, "夹具没有「样例表含占位符、被清」的块，判据空转"
    assert set(meta.rendered_block_hashes) == set(delivered)
    assert {c: h for c, h in meta.rendered_block_hashes.items() if delivered[c] != h} == {}


def test_nested_marker_paragraph_with_bookmark_is_cleared_not_deleted():
    """内容控件里带书签的标记段落只清空（书签不能跟着删）；没有书签的照旧删；body 级开闭标记照旧删。"""
    from docx.oxml import OxmlElement
    from docx.oxml.ns import qn

    from app.services.content_control_injector import inject_content_controls_for_blocks
    from app.services.word_doc_utils import remove_section_markers, scan_section_blocks

    doc = Document()
    doc.add_paragraph("##SECTION:八、1##")
    marked = doc.add_paragraph("##STYLE_REF:table:八、1:0##")
    doc.add_paragraph("##STYLE_REF:table:八、1:1##")
    doc.add_paragraph("正文")
    doc.add_paragraph("##/SECTION:八、1##")
    bookmark = OxmlElement("w:bookmarkStart")
    bookmark.set(qn("w:id"), "7")
    bookmark.set(qn("w:name"), "ref_x")
    marked._p.insert(0, bookmark)
    assert inject_content_controls_for_blocks(doc, scan_section_blocks(doc)) == 1
    remove_section_markers(doc)
    body = doc.element.body
    assert body.find(f".//{W}bookmarkStart[@{W}name='ref_x']") is not None
    paragraphs = ["".join(t.text or "" for t in p.iter(W + "t")) for p in body.iter(W + "p")]
    assert paragraphs == ["", "正文"], "带书签的标记段落应只清空，其余标记段落删除"


@pytest.mark.skipif(not _REAL_TEMPLATE.exists(), reason="真实附注模板缺失")
@pytest.mark.asyncio
async def test_real_template_fills_truncated_section_and_leaves_no_markers(monkeypatch):
    from app.services.note_section_catalog import build_variant_key

    index = {e["section_code"]: e for e in nwe_mod._load_section_code_index(build_variant_key("soe", "standalone"))}
    assert TRUNC in index and index[TRUNC].get("section_id"), "真实索引前提变了，判据空转"
    monkeypatch.setattr(nwe_mod.settings, "DELIVERABLE_LINEAGE_CONTENT_CONTROL_ENABLED", True)
    data = await _export([_note(FULL, "本公司以权责发生制为记账基础（第26组）。", index[TRUNC]["section_id"])])
    texts = _texts(data)
    assert "本公司以权责发生制为记账基础（第26组）。" in texts
    assert [t for t in texts if _LEFTOVER.search(t)][:5] == []


#: 测试集团甲实测章节码：锚点 sec_四_优先股_永续债等 逆映射成「四、优先股_永续债等」
#: （夹具 _patch_scope 把「九、」开头的章节当合并口径排除，故不用安徽的「九、1、或有负债」）
MULTI = "四、优先股、永续债等"


@pytest.mark.parametrize("controls", [True, False], ids=["内容控件开", "内容控件关"])
@pytest.mark.asyncio
async def test_rendered_block_hash_recomputed_for_multi_dun_section_code(tmp_path, monkeypatch, controls):
    """7.6 重算按真实章节码取块（spec deliverable-trio-generation-accuracy 16b 追加）。

    原实现按逆映射出的章节码查 ``rendered_block_hashes`` ⇒ 多「、」章节（安徽 2 节、测试集团甲 3 节）的哈希停在 6.4，
    收尾清理改了块内文字时，没动过的块也被增量刷新当成「有人工修改」。现行模板重建在 6.4 前已清掉块内占位符（真实
    交付件两份实测 6.4 与终稿一致），故用 7.5 包一层改块内文字来模拟「收尾清理改了这一块」。
    """
    from docx.oxml import OxmlElement

    from app.services.section_anchor_utils import block_text_hash, resolve_section_blocks

    body_text = "本公司未发行优先股、永续债等金融工具。"
    original_strip = nwe_mod._strip_residual_note_placeholders

    def _strip_and_touch(document):
        leaked = original_strip(document)
        for p in document.element.body.iter(W + "p"):
            if "".join(t.text or "" for t in p.iter(W + "t")) == body_text:
                run, text = OxmlElement("w:r"), OxmlElement("w:t")
                text.text = "（收尾清理后）"
                run.append(text)
                p.append(run)
        return leaked

    monkeypatch.setattr(nwe_mod, "_strip_residual_note_placeholders", _strip_and_touch)

    doc = Document()
    for code, body in ((MULTI, ["优先股、永续债等", "{{section:%s}}" % MULTI, "{{section:八、99}}"]),
                       ("八、1", ["货币资金", "{{section:八、1}}"])):
        doc.add_paragraph(f"##SECTION:{code}##")
        for line in body:
            doc.add_paragraph(line)
        doc.add_paragraph(f"##/SECTION:{code}##")
    path = tmp_path / "soe_standalone.docx"
    doc.save(str(path))
    _patch_loader(monkeypatch, path)
    _patch_index(monkeypatch, {MULTI: "sid-perpetual"})
    _patch_scope(monkeypatch)
    monkeypatch.setattr(nwe_mod.settings, "DELIVERABLE_LINEAGE_CONTENT_CONTROL_ENABLED", controls)

    exporter = NoteWordExporter(_make_db([_note(MULTI, body_text, "sid-perpetual"),
                                          _note("八、1", "货币资金期末余额披露", "sid-cash")]))
    buf, meta = await exporter.export_with_meta(
        nwe_mod.UUID(int=1), 2025, template_type="soe", report_scope="standalone", mode="template")
    data = buf.getvalue()
    assert body_text + "（收尾清理后）" in _texts(data), "夹具：收尾清理没改到这一块，判据空转"
    assert not any("{{section:八、99}}" in t for t in _texts(data))
    mode, blocks = resolve_section_blocks(Document(BytesIO(data)), known_codes=meta.kept_codes)
    delivered = {b.section_code: block_text_hash(list(b.elements)) for b in blocks}
    assert mode == "anchor" and MULTI in delivered and MULTI in meta.rendered_block_hashes
    assert meta.rendered_block_hashes[MULTI] == delivered[MULTI]

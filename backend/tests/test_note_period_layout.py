"""附注表格期间层（spec note-pull-header-formula-closure R4 / ADR-3）：只在显示 / 导出还原，不改存储。

判据：①布局文件与底稿 / 模板配对结果一致（生成器 --check 同口径）②读时按表名挂 ``_display_layout``，同组缺表不挂
③按期间拆开的两张表合成一张：标签列一份、数据列依次排开、表头外包期间；逐行行名不一致不合成 ④Word 程序化导出
画出 N 行表头（期间 / 账面余额·坏账准备 / 金额·比例），单元格值落在对的列。
"""
from __future__ import annotations

import copy

import pytest
from docx import Document

from app.services import note_period_layout as npl

END = {
    "name": "（2）按坏账准备计提方法分类披露应收账款",
    "headers": ["类 别", "金额", "比例(%)", "金额", "预期信用损失率(%)", "账面价值"],
    "_column_groups": [{"group": "账面余额", "start": 1, "span": 2}, {"group": "坏账准备", "start": 3, "span": 2}],
    "rows": [{"label": "按单项计提坏账准备", "values": [100.0, 10.0, 20.0, 20.0, 80.0], "_cell_modes": {"4": "manual"}},
             {"label": "合 计", "values": [1000.0, 100.0, 65.0, 6.5, 935.0], "is_total": True}],
}
BEGIN = copy.deepcopy(END)
BEGIN["name"] = "（2）按坏账准备计提方法分类披露应收账款（续：期初数）"
BEGIN["rows"][0]["values"] = [90.0, 9.0, 18.0, 20.0, 72.0]
BEGIN["rows"][0]["_cell_modes"] = {"0": "locked"}


def test_layout_file_is_current():
    """布局文件 = 现跑生成器的结果（底稿 xlsx / 单体模板 / 白名单改了要重跑 gen_note_table_period_layout.py）。"""
    assert npl.LAYOUT_PATH.read_text(encoding="utf-8") == npl.render_layout_document(npl.build_layout_document())


def test_known_layouts_from_workpapers():
    soe = npl.load_layouts()["soe"]["八、5"]
    side = next(e for e in soe if e["mode"] == "side_by_side")
    assert side["tables"] == [END["name"], BEGIN["name"]] and side["periods"] == ["期末数", "期初数"]
    assert {"mode": "top", "periods": ["期末数"]}.items() <= next(
        e for e in soe if e["tables"] == ["组合计提项目：应收中央企业客户"]).items()


def test_attach_by_table_name_and_whole_group_only():
    td = {"_tables": [copy.deepcopy(END), copy.deepcopy(BEGIN), {"name": "组合计提项目：应收中央企业客户"}]}
    assert npl.attach_display_layout(td, "soe", "八、5") == 3
    assert td["_tables"][0]["_display_layout"] == {"mode": "side_by_side", "period": "期末数", "group": END["name"],
                                                   "order": 0}
    assert td["_tables"][1]["_display_layout"]["order"] == 1
    assert td["_tables"][2]["_display_layout"] == {"mode": "top", "period": "期末数"}
    lonely = {"_tables": [copy.deepcopy(END)]}
    assert npl.attach_display_layout(lonely, "soe", "八、5") == 0 and "_display_layout" not in lonely["_tables"][0]
    other_std = {"_tables": [copy.deepcopy(END), copy.deepcopy(BEGIN)]}
    assert npl.attach_display_layout(other_std, "listed", "八、5") == 0, "上市没有这一节的布局：口径不串"


def test_header_defs_three_levels():
    paths = [("类别",), ("期末数", "账面余额", "金额"), ("期末数", "账面余额", "比例"), ("期末数", "账面价值"),
             ("期初数", "账面余额", "金额")]
    rows = npl.header_defs(paths)
    assert rows == [
        [{"text": "类别", "colspan": 1, "rowspan": 3}, {"text": "期末数", "colspan": 3, "rowspan": 1},
         {"text": "期初数", "colspan": 1, "rowspan": 1}],
        [{"text": "账面余额", "colspan": 2, "rowspan": 1}, {"text": "账面价值", "colspan": 1, "rowspan": 2},
         {"text": "账面余额", "colspan": 1, "rowspan": 1}],
        [{"text": "金额", "colspan": 1, "rowspan": 1}, {"text": "比例", "colspan": 1, "rowspan": 1},
         {"text": "金额", "colspan": 1, "rowspan": 1}],
    ]


def test_side_by_side_merges_and_keeps_coordinates():
    items = [(1, copy.deepcopy(END), False), (2, copy.deepcopy(BEGIN), True), (3, {"name": "别的表"}, False)]
    out = npl.apply_period_layout_indexed(items, "soe", "八、5")
    assert [i for i, _t, _e in out] == [1, 3] and out[0][2] is False
    merged = out[0][1]
    assert merged["headers"][0] == "类 别" and len(merged["headers"]) == 11
    assert merged["rows"][0]["values"] == [100.0, 10.0, 20.0, 20.0, 80.0, 90.0, 9.0, 18.0, 20.0, 72.0]
    assert merged["rows"][0]["_cell_modes"] == {"4": "manual", "5": "locked"}, "单元格模式按列偏移带过来"
    assert merged["_header_paths"][1] == ["期末数", "账面余额", "金额"]
    assert merged["_header_paths"][10] == ["期初数", "账面价值"]
    assert "_column_groups" not in merged and merged["_merged_from"] == [END["name"], BEGIN["name"]]
    assert "_display_layout" not in items[0][1], "不改入参"


def test_side_by_side_needs_identical_row_labels():
    other = copy.deepcopy(BEGIN)
    other["rows"][0]["label"] = "按组合计提坏账准备"
    out = npl.apply_period_layout_indexed([(0, copy.deepcopy(END), False), (1, other, False)], "soe", "八、5")
    assert [t["name"] for _i, t, _e in out] == [END["name"], BEGIN["name"]], "行名不一致不合成"
    assert all("_header_paths" not in t and "_display_layout" not in t for _i, t, _e in out)


def test_top_adds_period_layer():
    table = {"name": "组合计提项目：应收中央企业客户", "headers": ["账 龄", "应收账款", "比例（%）", "坏账准备"],
             "rows": [{"label": "1年以内", "values": [1, 2, 3]}]}
    (_i, laid, _e), = npl.apply_period_layout_indexed([(4, table, False)], "soe", "八、5")
    assert laid["_header_paths"] == [["账 龄"], ["期末数", "应收账款"], ["期末数", "比例（%）"], ["期末数", "坏账准备"]]
    assert npl.export_header_defs(laid)[0] == [{"text": "账 龄", "colspan": 1, "rowspan": 2},
                                               {"text": "期末数", "colspan": 3, "rowspan": 1}]


def test_two_level_groups_keep_existing_export_path():
    assert npl.export_header_defs(copy.deepcopy(END)) is None, "两层分组照旧由导出器画两行"


def _row_texts(table, r):
    """一行里各（合并后）单元格的字：持有元素引用去重（lxml 代理的 id 会被复用，不能按 id 去重）。"""
    kept, out = [], []
    for cell in table.rows[r].cells:
        if not any(cell._tc is k for k in kept):
            kept.append(cell._tc)
            out.append(cell.text)
    return out


def _row_height(row):
    """行高（致同表样：每行 exact 0.7cm = 397 twip）。"""
    from docx.oxml.ns import qn

    tr_pr = row._tr.find(qn("w:trPr"))
    h = tr_pr.find(qn("w:trHeight")) if tr_pr is not None else None
    return None if h is None else (h.get(qn("w:val")), h.get(qn("w:hRule")))


def test_word_export_two_level_groups_draw_two_header_rows():
    """两层分组照旧画两行表头（与 N 行表头共用 ``fill_multi_header`` 一条路径）：分组名横跨子列、无分组列跨两行。"""
    from app.services.note_word_exporter import NoteWordExporter

    doc = Document()
    NoteWordExporter(None)._render_table(doc, copy.deepcopy(END))  # type: ignore[arg-type]
    table = doc.tables[-1]
    assert len(table.columns) == 6 and len(table.rows) == 2 + len(END["rows"])
    assert _row_texts(table, 0) == ["类 别", "账面余额", "坏账准备", "账面价值"]
    assert _row_texts(table, 1)[1:5] == ["金额", "比例(%)", "金额", "预期信用损失率(%)"]
    assert table.rows[0].cells[0]._tc is table.rows[1].cells[0]._tc, "无分组的标签列跨两行"
    assert [_row_height(table.rows[r]) for r in range(2)] == [("397", "exact")] * 2, "两行表头都按表样定行高"
    assert table.rows[2].cells[1].text.replace(",", "") in ("100.00", "100")
    assert table.rows[2].cells[5].text.replace(",", "") in ("80.00", "80")


def test_word_export_draws_three_header_rows():
    from app.services.note_word_exporter import NoteWordExporter

    merged = npl.apply_period_layout([copy.deepcopy(END), copy.deepcopy(BEGIN)], "soe", "八、5")[0]
    doc = Document()
    NoteWordExporter(None)._render_table(doc, merged)  # type: ignore[arg-type]
    table = doc.tables[-1]
    assert len(table.columns) == 11 and len(table.rows) == 3 + 2

    def texts(r):
        return _row_texts(table, r)

    assert texts(0) == ["类 别", "期末数", "期初数"]
    assert texts(1)[1:] == ["账面余额", "坏账准备", "账面价值", "账面余额", "坏账准备", "账面价值"]
    assert texts(2)[1:5] == ["金额", "比例(%)", "金额", "预期信用损失率(%)"]
    assert [_row_height(table.rows[r]) for r in range(3)] == [("397", "exact")] * 3, "三行表头都按表样定行高"
    assert table.rows[0].cells[0]._tc is table.rows[2].cells[0]._tc, "标签列跨三行"
    assert table.rows[3].cells[1].text.replace(",", "") in ("100.00", "100")
    assert table.rows[3].cells[6].text.replace(",", "") in ("90.00", "90")


def test_exporters_table_list_applies_layout():
    """程序化导出与模板模式取表都走期间层：按期间拆开的两张合成一张（模板模式取首张的表序，可对上 Word 样例表）。"""
    from types import SimpleNamespace

    from app.services.note_template_block_fill import note_tables_indexed
    from app.services.note_word_exporter import NoteWordExporter

    note = SimpleNamespace(note_section="八、5", source_template="soe",
                           table_data={"_tables": [copy.deepcopy(END), copy.deepcopy(BEGIN)]})
    exporter = NoteWordExporter(None)  # type: ignore[arg-type]
    tables = exporter._note_tables(note)
    assert len(tables) == 1 and tables[0]["_merged_from"] == [END["name"], BEGIN["name"]]
    indexed = note_tables_indexed(exporter, note)
    assert [(i, t.get("_merged_from")) for i, t, _e in indexed] == [(0, [END["name"], BEGIN["name"]])]


def test_read_time_enrichment_attaches_and_replaces_stale_marks():
    """详情端点读时增强挂布局；前端保存会把读时的键带回库 ⇒ 每次以布局文件为准（存进去的旧标记先清掉）。"""
    from app.services.note_table_enrichment import enrich_note_table_data

    td = {"_tables": [copy.deepcopy(END), copy.deepcopy(BEGIN), {"name": "别的表", "_display_layout": {"mode": "top"}}]}
    enrich_note_table_data(td, "soe", "八、5")
    assert td["_tables"][0]["_display_layout"]["mode"] == "side_by_side"
    assert "_display_layout" not in td["_tables"][2], "布局文件里没有的表：旧标记清掉"


@pytest.mark.parametrize("bad", [None, {}, {"_tables": "x"}])
def test_attach_tolerates_bad_input(bad):
    assert npl.attach_display_layout(bad, "soe", "八、5") == 0

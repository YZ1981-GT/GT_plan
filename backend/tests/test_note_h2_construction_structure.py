"""附注在建工程章节结构守卫（§五、23 上市 / §八、23 国企）。

锁定 `fix_note_h2_construction_structure.py` 的对齐结果：两级表头（在建工程明细/国企汇总/
国企情况）、单级 flat 表、表名（工程物资 非「项  目」）、columns key 与前端载荷一致、guidance。
含 openpyxl 直读源 xlsx tab 名 + 两级表头交叉比对 + 反向自检。

spec: h2-construction-in-progress-disclosure-alignment
"""
from __future__ import annotations

import importlib.util
import json
from pathlib import Path

import openpyxl
import pytest

_ROOT = Path(__file__).resolve().parent.parent
_FIX = _ROOT / "scripts" / "fix" / "fix_note_h2_construction_structure.py"
_SRC_XLSX = _ROOT / "wp_templates" / "H" / "H2 在建工程.xlsx"


def _load_fix():
    spec = importlib.util.spec_from_file_location("_fix_h2_struct", _FIX)
    mod = importlib.util.module_from_spec(spec)
    spec.loader.exec_module(mod)
    return mod


FIX = _load_fix()

# 权威列 key（逐字取自 h2DisclosureSyncPayload.ts）
_EXPECTED_KEYS = {
    "listed": {
        "在建工程": ["label", "end_balance", "prior_balance"],
        "在建工程明细": [
            "label", "end_book", "end_impairment", "end_net",
            "prior_book", "prior_impairment", "prior_net",
        ],
        "重要在建工程项目变动情况": [
            "label", "begin_balance", "increase", "transfer_to_fa", "other_decrease",
            "interest_cap_accum", "interest_cap_current", "interest_cap_rate", "end_balance",
        ],
        "重要在建工程项目变动情况（续）：": ["label", "budget", "cum_input_pct", "progress", "fund_source"],
        "在建工程减值准备情况": ["label", "begin_balance", "provision", "decrease", "end_balance"],
        "工程物资": ["label", "end_balance", "prior_balance"],
    },
    "soe": {
        "在建工程": [
            "label", "end_book", "end_impairment", "end_carrying",
            "begin_book", "begin_impairment", "begin_carrying",
        ],
        "（1）在建工程情况": [
            "label", "end_book", "end_impairment", "end_carrying",
            "begin_book", "begin_impairment", "begin_carrying",
        ],
        "（2）重要在建工程项目本期变动情况": [
            "label", "budget", "begin_balance", "increase", "transfer_to_fa", "other_decrease",
            "end_balance", "cum_input_pct", "progress", "interest_cap_accum",
            "interest_cap_current", "interest_cap_rate", "fund_source",
        ],
        "（3）本期计提在建工程减值准备情况": ["label", "provision_amount", "reason"],
    },
}

# 两级表（应有 _column_groups 2 组）
_TWO_LEVEL = {
    "listed": {"在建工程明细"},
    "soe": {"在建工程", "（1）在建工程情况"},
}


def _section(variant: str) -> dict:
    path, section_number, _ = FIX._TARGETS[variant]
    doc = json.loads(Path(path).read_text(encoding="utf-8"))
    return next(s for s in doc["sections"] if str(s.get("section_number")) == section_number)


@pytest.mark.parametrize("variant", ["listed", "soe"])
def test_check_passes(variant: str):
    _changes, warnings, errs = FIX._runner(variant, dry_run=True, check=True)
    assert not errs, f"{variant} 结构欠账：{errs}"
    assert not warnings, f"{variant} 告警：{warnings}"


@pytest.mark.parametrize("variant", ["listed", "soe"])
def test_table_names(variant: str):
    names = [str(t.get("name")) for t in _section(variant).get("tables") or []]
    assert names == FIX.EXPECTED[variant], f"{variant} 表序/表名漂移：{names}"
    assert "项  目" not in names, f"{variant} 残留垃圾表名「项  目」"


@pytest.mark.parametrize("variant", ["listed", "soe"])
def test_columns_and_groups(variant: str):
    by_name = {str(t.get("name")): t for t in _section(variant).get("tables") or []}
    for name, want_keys in _EXPECTED_KEYS[variant].items():
        tbl = by_name[name]
        cols = tbl.get("columns") or []
        assert [c.get("key") for c in cols] == want_keys, f"{variant}/{name} 列 key 漂移"
        assert cols[0].get("is_label"), f"{variant}/{name} 首列未 is_label"
        assert str(cols[0].get("label")) == str(tbl["headers"][0]), f"{variant}/{name} 首列≠headers[0]"
        assert str(tbl.get("guidance") or "").strip(), f"{variant}/{name} 缺 guidance"
        assert not any(r.get("row_type") == "header_label" for r in tbl.get("rows") or []), (
            f"{variant}/{name} 残留 header_label 假行"
        )
        if name in _TWO_LEVEL[variant]:
            groups = tbl.get("_column_groups") or []
            assert len(groups) == 2, f"{variant}/{name} 两级表头应有 2 分组，实得 {groups}"
            assert all(c.get("group") for c in cols[1:]), f"{variant}/{name} 数据列缺 group"
        else:
            assert not tbl.get("_column_groups"), f"{variant}/{name} 单级表残留 _column_groups"
            assert any(c.get("flat") for c in cols), f"{variant}/{name} 单级表未标 flat"


def test_source_xlsx_two_level_headers():
    """openpyxl 直读源 xlsx：两级表头叶子名与模板一致（防结构确权漂移）。"""
    wb = openpyxl.load_workbook(_SRC_XLSX, data_only=True)
    listed = wb["附注披露信息（上市公司）"]
    # 上市在建工程明细：行14 = ['', 账面余额, 减值准备, 账面净值, 账面余额, 减值准备, 账面净值]
    row14 = [listed.cell(14, c).value for c in range(2, 8)]
    assert row14 == ["账面余额", "减值准备", "账面净值", "账面余额", "减值准备", "账面净值"], row14
    soe = wb["附注披露信息（国有企业）"]
    # 国企汇总：行8 = ['', 账面余额, 减值准备, 账面价值, 账面余额, 减值准备, 账面价值]
    row8 = [soe.cell(8, c).value for c in range(2, 8)]
    assert row8 == ["账面余额", "减值准备", "账面价值", "账面余额", "减值准备", "账面价值"], row8


def test_reverse_self_check():
    """validate_section 能抓出缺 columns / header_label 残留（防守卫空转）。"""
    from _note_structure_kit import validate_section

    broken = {
        "tables": [{
            "name": "在建工程",
            "headers": ["项目", "期末余额"],
            "rows": [{"label": "项目", "row_type": "header_label"}],
            "guidance": "x",
        }],
    }
    errs = validate_section(broken, ["在建工程"])
    assert any("缺 columns" in e for e in errs), errs
    assert any("header_label" in e for e in errs), errs

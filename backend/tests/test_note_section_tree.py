"""第 23 组：上传文档按章节树识别（AC-1.6）与目标树的建立。

用户拍板（2026-10-15）：「识别上传文档时，也需要按章节树形来」。首版 PDF 只分 1 / 2 两级、逐 span 判标题，
中文章节号纯文字一律 2 级 ⇒ 都不成树。
"""

from __future__ import annotations

import uuid

import pytest

from app.services.note_document_parser import ParsedSection, ParseResult, heading_rank, parse_pdf, parse_text_source
from app.services.note_section_tree_locator import build_source_tree, build_target_tree, source_parent_ids
from app.services.note_text_batch_service import tree_warnings


def _s(sid: str, title: str, level: int, doc: str | None = None) -> ParsedSection:
    return ParsedSection(source_id=sid, title=title, level=level, metadata={"doc_id": doc} if doc else {})


@pytest.mark.parametrize(("text", "rank"), [
    ("一、公司基本情况", 1), ("十一、关联方及关联交易", 1), ("（一）主要税种及税率", 2), ("(二)税收优惠", 2),
    ("1.会计期间", 3), ("12、其他", 3), ("1．记账本位币", 3), ("（1）共同经营", 4), ("(2)合营企业", 4), ("①消耗性生物资产", 5),
    ("X、数据资源", 3),  # 国企模板「项目注释」里编号待定的节写作「X、」
    ("1.5亿元的借款", None), ("会计期间", None), ("", None), ("X.509证书", None),
])
def test_heading_rank(text, rank):
    assert heading_rank(text) == rank


class TestSourceTree:
    def test_parent_is_nearest_higher_level_heading(self):
        """国企 Word：章 H1、节 H3、节内小标题 H4 —— 层级可跳，上级仍是前面最近的更高层级标题。"""
        sections = [_s("c", "重要会计政策、会计估计", 1), _s("a", "金融工具", 3), _s("b", "金融资产分类", 4),
                    _s("d", "存货", 3), _s("e", "财务报表主要项目注释", 1)]
        assert source_parent_ids(sections) == {"c": None, "a": "c", "b": "a", "d": "c", "e": None}

    def test_untitled_placeholder_is_never_a_parent(self):
        sections = [_s("u", "（无标题）", 1), _s("a", "会计期间", 2), _s("b", "记账本位币", 2)]
        assert source_parent_ids(sections) == {"u": None, "a": None, "b": None}

    def test_each_knowledge_document_is_its_own_tree(self):
        sections = [_s("d1:s1", "公司基本情况", 1, "d1"), _s("d2:s1", "补充说明", 2, "d2")]
        assert source_parent_ids(sections) == {"d1:s1": None, "d2:s1": None}
        assert [len(n.children) for n in build_source_tree(sections)] == [0, 0]


class TestTextSourcesBuildTree:
    def test_cn_numbered_plain_text_levels_by_rank(self):
        text = "一、重要会计政策\n（一）会计期间\n自公历1月1日起。\n1. 记账本位币\n人民币。\n二、税项\n增值税。"
        sections = parse_text_source(text).sections
        assert [(s.title, s.level) for s in sections] == [
            ("一、重要会计政策", 1), ("（一）会计期间", 2), ("1. 记账本位币", 3), ("二、税项", 1)]
        parents = source_parent_ids(sections)
        by_title = {s.title: s.source_id for s in sections}
        assert parents[by_title["1. 记账本位币"]] == by_title["（一）会计期间"]
        assert parents[by_title["二、税项"]] is None

    def test_pdf_line_with_two_spans_is_one_heading_and_ranks_build_tree(self):
        """Word 自动编号导出 PDF 后编号与标题是同一行的两个 span：首版逐 span 判，「1.」成标题、「会计期间」成正文。"""
        import pymupdf

        doc = pymupdf.open()
        page = doc.new_page()
        writer = pymupdf.TextWriter(page.rect)
        chinese, latin = pymupdf.Font("china-s"), pymupdf.Font("helv")
        writer.append((72, 72), "一、重要会计政策", font=chinese, fontsize=18)
        writer.append((72, 110), "1.", font=latin, fontsize=16)
        writer.append(writer.last_point, "会计期间", font=chinese, fontsize=16)
        writer.append((72, 140), "本公司会计期间为公历1月1日至12月31日，本报告期为2025年度。", font=chinese, fontsize=10)
        writer.append((72, 160), "本公司以人民币为记账本位币，境外子公司按所在地货币记账。", font=chinese, fontsize=10)
        writer.write_text(page)
        data = doc.tobytes()
        lines = [ln for b in page.get_text("dict")["blocks"] if b.get("type") == 0 for ln in b["lines"]]
        doc.close()
        assert any(len([s for s in ln["spans"] if s["text"].strip()]) >= 2 for ln in lines), "夹具未产生同行两个 span"

        sections = parse_pdf(data).sections
        assert [(s.title, s.level) for s in sections] == [("一、重要会计政策", 1), ("1.会计期间", 3)]
        assert source_parent_ids(sections)[sections[1].source_id] == sections[0].source_id
        assert "本公司会计期间为公历" in sections[1].html


def _outline_pdf() -> bytes:
    """标题与正文同字号、不加粗（LibreOffice 把国企单体 Word 转 PDF 即此形态：字号判标题只得 3 章），但带书签目录。

    书签里混进两条表格单元格（「账 龄」「期末数」：模板表头段落误设了 1 级大纲，真实 PDF 第 179 页同形，全文 39 条），
    书签是一棵树、深度相对上级算 ⇒ 其后的「（2）」「X、」深度被算成 2（真实层级 3 / 2）；
    另有一条跨两行的长标题、一条编号待定的「X、」节、一条页上找不到的书签。
    """
    import pymupdf

    doc = pymupdf.open()
    font = pymupdf.Font("china-s")
    pages = [
        [(72, 60, "财务报表附注"), (72, 90, "一、重要会计政策"), (72, 110, "本公司执行企业会计准则。"),
         (72, 130, "1、会计期间"), (72, 150, "本公司会计期间为公历1月1日至12月31日。")],
        [(72, 60, "2、应收账款"), (72, 80, "（1）按欠款方归集的期末余额前五名"), (72, 96, "的应收账款情况"),
         (72, 116, "前五名合计占比35%。"), (72, 136, "账 龄"), (300, 156, "期末数"),
         (72, 176, "（2）坏账准备计提情况"), (72, 196, "本期计提坏账准备100万元。"),
         (72, 216, "①单项计提坏账准备"), (72, 236, "单项计提40万元。"),
         (72, 256, "X、数据资源"), (72, 276, "本期无数据资源。"), (72, 296, "二、税项"), (72, 316, "增值税税率13%。")],
    ]
    for lines in pages:
        page = doc.new_page()
        writer = pymupdf.TextWriter(page.rect)
        for x, y, text in lines:
            writer.append((x, y), text, font=font, fontsize=11)
        writer.write_text(page)
    doc.set_toc([  # 「2、应收账款」书签指向第 1 页（段前间距落在上一页页底，标题文字在下一页页首）
        [1, "一、 重要会计政策", 1], [2, "1、 会计期间", 1], [2, "2、 应收账款", 1],
        [3, "（1）按欠款方归集的期末余额前五名的应收账款情况", 2], [1, "账 龄", 2], [1, "期末数", 2],
        [2, "（2）坏账准备计提情况", 2], [3, "①单项计提坏账准备", 2], [2, "X、数据资源", 2],
        [2, "3、 不存在的标题", 2], [1, "二、 税项", 2],
    ])
    data = doc.tobytes()
    doc.close()
    return data


class TestPdfOutline:
    """PDF 带书签目录时按书签成树（第 23 组 23.5 实测：国企单体 PDF 标题与正文同为 12 号不加粗，字号判只得 3 章；书签 422 条层级齐全）。"""

    def test_sections_follow_outline_titles_levels_and_split_body(self):
        result = parse_pdf(_outline_pdf())
        assert [(s.title, s.level) for s in result.sections] == [
            ("（无标题）", 1), ("一、 重要会计政策", 1), ("1、 会计期间", 2), ("2、 应收账款", 2),
            ("（1）按欠款方归集的期末余额前五名的应收账款情况", 3), ("（2）坏账准备计提情况", 3), ("①单项计提坏账准备", 4),
            ("X、数据资源", 2), ("二、 税项", 1)]
        html = {s.title: s.html for s in result.sections}
        assert "财务报表附注" in html["（无标题）"]
        assert "本公司执行企业会计准则" in html["一、 重要会计政策"]
        assert "本公司会计期间为公历" in html["1、 会计期间"]
        assert html["2、 应收账款"] == ""
        first = html["（1）按欠款方归集的期末余额前五名的应收账款情况"]
        assert "前五名合计占比" in first and "期末数" in first  # 表格单元格仍是正文
        assert "的应收账款情况" not in first  # 跨两行的标题整体是标题
        assert "本期计提坏账准备" in html["（2）坏账准备计提情况"] and "单项计提40万元" in html["①单项计提坏账准备"]
        assert "本期无数据资源" in html["X、数据资源"] and "增值税税率" in html["二、 税项"]

    def test_unnumbered_entries_beside_numbered_siblings_are_not_headings(self):
        """表格单元格书签若当标题：「（2）」「X、」挂到「期末数」下，整棵子树定位失败；去掉后深度须按编号对齐，
        只删不对齐则「（2）」深度 2 成了「一」的直接下级（与「（1）」不再同级）。"""
        result = parse_pdf(_outline_pdf())
        by_title = {s.title: s.source_id for s in result.sections}
        parents = source_parent_ids(result.sections)
        assert parents[by_title["2、 应收账款"]] == by_title["一、 重要会计政策"]
        assert parents[by_title["（1）按欠款方归集的期末余额前五名的应收账款情况"]] == by_title["2、 应收账款"]
        assert parents[by_title["（2）坏账准备计提情况"]] == by_title["2、 应收账款"]
        assert parents[by_title["①单项计提坏账准备"]] == by_title["（2）坏账准备计提情况"]  # 重排层级的子树整体平移
        assert parents[by_title["X、数据资源"]] == by_title["一、 重要会计政策"]
        assert parents[by_title["二、 税项"]] is None
        dropped = [w for w in result.warnings if "无编号" in w]
        assert len(dropped) == 1 and "2 条" in dropped[0] and "「账 龄」" in dropped[0] and "「期末数」" in dropped[0]

    def test_outline_entry_missing_on_its_page_is_reported_not_guessed(self):
        result = parse_pdf(_outline_pdf())
        assert "3、 不存在的标题" not in {s.title for s in result.sections}
        assert any("「3、 不存在的标题」" in w and "第 2 页" in w for w in result.warnings)

    def test_same_title_twice_on_one_page_is_located_in_document_order(self):
        """「（1）按账龄披露」在应收账款、其他应收款下各有一个，常在同一页：每条书签只能从上一条之后找。"""
        import pymupdf

        doc = pymupdf.open()
        page = doc.new_page()
        writer = pymupdf.TextWriter(page.rect)
        font = pymupdf.Font("china-s")
        for y, text in ((60, "1、应收账款"), (80, "（1）按账龄披露"), (100, "应收账款账龄正文。"),
                        (120, "2、其他应收款"), (140, "（1）按账龄披露"), (160, "其他应收款账龄正文。")):
            writer.append((72, y), text, font=font, fontsize=11)
        writer.write_text(page)
        doc.set_toc([[1, "1、应收账款", 1], [2, "（1）按账龄披露", 1], [1, "2、其他应收款", 1], [2, "（1）按账龄披露", 1]])
        data = doc.tobytes()
        doc.close()

        sections = parse_pdf(data).sections
        assert [(s.title, s.level) for s in sections] == [
            ("1、应收账款", 1), ("（1）按账龄披露", 2), ("2、其他应收款", 1), ("（1）按账龄披露", 2)]
        assert [s.html for s in sections] == ["", "<p>应收账款账龄正文。</p>", "", "<p>其他应收款账龄正文。</p>"]

    @pytest.mark.parametrize("with_outline", [True, False])
    def test_running_header_and_page_number_are_not_body(self, with_outline):
        """每页页眉（公司名 / 财务报表附注 / 年度）与页脚页码首版混进章节正文：跨页的节中间插一段页眉，
        章标题在页底、首个下级在下页页首时，章正文只剩页眉页脚 ⇒ 多出一条「章首正文」要用户手工处理
        （第 23 组国企单体 PDF 实测「四、重要会计政策」）。"""
        import pymupdf

        doc = pymupdf.open()
        font = pymupdf.Font("china-s")
        unit = (300, "单位：元", 11)  # 每页正文中部都有：不在页眉页脚边带里，是正文
        bodies = [[(90, "一、重要会计政策", 16), (110, "本公司执行企业会计准则。", 11), unit, (780, "二、税项", 16)],
                  [(90, "1、主要税种", 15), (110, "增值税税率13%。", 11), unit],
                  [(90, "企业所得税税率25%。", 11), unit],
                  [(90, "三、其他事项", 16), (110, "无其他事项。", 11), unit]]
        for number, lines in enumerate(bodies, start=1):
            page = doc.new_page()
            writer = pymupdf.TextWriter(page.rect)
            for y, text in ((40, "甲公司"), (56, "财务报表附注"), (812, str(number))):
                writer.append((72, y), text, font=font, fontsize=9)
            for y, text, size in lines:
                writer.append((72, y), text, font=font, fontsize=size)
            writer.write_text(page)
        if with_outline:
            doc.set_toc([[1, "一、重要会计政策", 1], [1, "二、税项", 1], [2, "1、主要税种", 2], [1, "三、其他事项", 4]])
        data = doc.tobytes()
        doc.close()

        sections = parse_pdf(data).sections
        assert [s.title for s in sections] == ["一、重要会计政策", "二、税项", "1、主要税种", "三、其他事项"]
        assert [s.html for s in sections] == [
            "<p>本公司执行企业会计准则。</p><p>单位：元</p>", "",
            "<p>增值税税率13%。</p><p>单位：元</p><p>企业所得税税率25%。</p><p>单位：元</p>",
            "<p>无其他事项。</p><p>单位：元</p>"]

    def test_entry_without_page_in_this_document_is_reported(self):
        from app.services.note_pdf_outline import OutlineHeading, locate_heading_lines

        headings = [OutlineHeading(0, "1、会计期间", -1, 1, 3), OutlineHeading(1, "2、记账本位币", 9, 1, 3)]
        hits, warnings = locate_heading_lines(["1、会计期间", "2、记账本位币"], [0, 2], headings)
        assert hits == []
        assert warnings == ["书签「1、会计期间」没有指向本文档的页面，未作为标题",
                            "书签「2、记账本位币」没有指向本文档的页面，未作为标题"]


def test_circled_number_is_a_numbering_prefix_for_matching():
    """Word 自动编号的「①承租人信息」：DOCX 文字层没有编号、PDF 书签 / 文字层带着 ⇒ 同一份附注两种格式规范标题须一致
    （四份模板转 PDF 实测：两种格式的树只差在这类标题上）。"""
    from app.services.note_section_kind import canonical_title

    assert canonical_title("①承租人信息") == canonical_title("承租人信息") == "承租人信息"
    assert canonical_title("⑳其他") == "其他"


def test_docx_style_inheriting_heading_outline_is_a_heading():
    """国企单体 Word「同一控制下的企业合并」用样式「样式2」（基于 Heading 4，大纲级别继承自基样式）：
    Word 导航窗格与 LibreOffice 书签都当 4 级标题，首版只认样式名与段落自带大纲级别 ⇒ 并进上一节正文。"""
    import io

    import docx
    from docx.enum.style import WD_STYLE_TYPE
    from docx.oxml import OxmlElement
    from docx.oxml.ns import qn

    from app.services.note_document_parser import parse_docx

    def body_level(element) -> None:
        outline = OxmlElement("w:outlineLvl")
        outline.set(qn("w:val"), "9")
        element.get_or_add_pPr().append(outline)

    document = docx.Document()
    custom = document.styles.add_style("样式2", WD_STYLE_TYPE.PARAGRAPH)
    custom.base_style = document.styles["Heading 4"]
    base_ppr = document.styles["Heading 4"].element.find(qn("w:pPr"))
    assert base_ppr is not None and base_ppr.find(qn("w:outlineLvl")) is not None, "夹具：Heading 4 须自带大纲级别"
    note_style = document.styles.add_style("说明文字", WD_STYLE_TYPE.PARAGRAPH)  # 基于标题、自身改回正文级别
    note_style.base_style = custom
    body_level(note_style.element)
    document.add_paragraph("重要会计政策", style="Heading 1")
    document.add_paragraph("企业合并", style="Heading 3")
    document.add_paragraph("同一控制下的企业合并", style="样式2")
    document.add_paragraph("参与合并的企业在合并前后均受同一方最终控制。")
    body_level(document.add_paragraph("样式2 但段落自设为正文级别", style="样式2")._p)
    document.add_paragraph("样式自设为正文级别的说明", style="说明文字")
    buffer = io.BytesIO()
    document.save(buffer)

    sections = parse_docx(buffer.getvalue()).sections
    assert [(s.title, s.level) for s in sections] == [("重要会计政策", 1), ("企业合并", 3), ("同一控制下的企业合并", 4)]
    assert all(text in sections[2].html for text in ("均受同一方最终控制", "段落自设为正文级别", "样式自设为正文级别的说明"))


# ── 知识库多篇文档各自成树（第 23 组，自 test_note_text_batch_endpoints.py 搬来：纯函数，不经端点）──


def _kb(sid: str, title: str, level: int, doc: str) -> object:
    return ParsedSection(source_id=sid, title=title, level=level, html=f"<p>{title}正文</p>", metadata={"doc_id": doc})


def _company_tree() -> tuple[list[dict], dict]:
    chapter = {"id": uuid.uuid4(), "note_section": "一", "section_title": "公司基本情况", "section_kind": "container",
               "section_id": "ch-01", "parent_section_id": None, "chapter_title": None}
    leaf = {"id": uuid.uuid4(), "note_section": "一、1", "section_title": "公司基本情况", "section_kind": "narrative",
            "section_id": "s-01-01", "parent_section_id": "ch-01", "chapter_title": "公司基本情况"}
    return [chapter, leaf], leaf


def test_knowledge_documents_do_not_fold_into_each_other():
    """多篇知识库文档拼接：各自成树，下一篇开头的深层标题不得成为上一篇末节的下级（被并入或在其下定位）。"""
    from app.services.note_text_batch_service import SourceContext, locate_sources

    targets, leaf = _company_tree()
    sources = [_kb("d1:s1", "公司基本情况", 1, "d1"), _kb("d2:s1", "补充说明", 2, "d2")]
    folded_sources, matches, _s = locate_sources(sources, targets, source_context=SourceContext(kind="knowledge"))
    by = {m.source_id: m for m in matches}
    assert by["d1:s1"].target_note_id == leaf["id"] and folded_sources[0].html == "<p>公司基本情况正文</p>"
    assert (by["d2:s1"].match_kind, by["d2:s1"].target_note_id) == ("tree_unmatched", None)


def test_knowledge_document_chapter_does_not_leak_into_next_document():
    """上一篇末尾的「项目注释」章不得把下一篇的章节当成其下级。"""
    from app.services.note_text_batch_service import SourceContext, locate_sources

    targets, leaf = _company_tree()
    sources = [_kb("d1:s1", "合并财务报表项目注释", 1, "d1"), _kb("d2:s1", "公司基本情况", 2, "d2")]
    by = {m.source_id: m for m in locate_sources(sources, targets, source_context=SourceContext(kind="knowledge"))[1]}
    assert by["d2:s1"].target_note_id == leaf["id"]


class TestTargetTree:
    @staticmethod
    def _row(code, title, sid=None, parent=None):
        return {"id": uuid.uuid4(), "note_section": code, "section_title": title, "section_id": sid,
                "parent_section_id": parent}

    def test_roots_children_and_off_tree_rows(self):
        rows = [self._row("一", "公司基本情况", "ch-01"), self._row("一、1", "公司基本情况", "s-1", "ch-01"),
                self._row("五、1", "货币资金", "listed-x", "chapter-05-he-bing"),  # 上级不存在（和平真库 40 行同形）
                self._row("四、记账基础和计价原", "记账基础和计价原则")]          # 无 section_id（真库 14~15 行同形）
        tree = build_target_tree(rows)
        assert [r["note_section"] for r in tree.roots] == ["一"]
        assert [r["note_section"] for r in tree.kids(rows[0])] == ["一、1"]
        assert {r["note_section"] for r in tree.off_tree} == {"五、1", "四、记账基础和计价原"}
        assert tree_warnings(rows) == ["本项目有 2 个章节不在章节树上（缺少层级信息或上级章节），不参与自动定位，可在下拉里手工指定"]

    def test_project_without_tree_is_told_plainly(self):
        assert "章节树未建立" in tree_warnings([self._row("一、1", "公司基本情况")])[0]
        assert tree_warnings([self._row("一", "公司基本情况", "ch-01")]) == []


def test_upload_payload_carries_parent_ids_for_the_tree_view():
    from app.services.note_text_batch_flow import upload_source_payload

    result = ParseResult(sections=[_s("docx-s1", "重要会计政策、会计估计", 1), _s("docx-s2", "会计期间", 3)], file_hash="h")
    payload = upload_source_payload("docx", "附注.docx", b"x", result, project_id=uuid.uuid4(), user_id=uuid.uuid4())
    assert [(s["source_id"], s["parent_source_id"]) for s in payload["sections"]] == [
        ("docx-s1", None), ("docx-s2", "docx-s1")]

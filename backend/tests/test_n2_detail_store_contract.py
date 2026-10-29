# -*- coding: utf-8 -*-
"""N2-2 前端、导入导出、复核与 workpaper-sync 的 store 契约闭合测试。"""
from __future__ import annotations

import io
import json
from types import SimpleNamespace
from typing import Any
from uuid import UUID

import openpyxl
import pytest
from fastapi import HTTPException

from app.services.n2_detail_model import (
    N2_DETAIL_FIELDS,
    N2_DETAIL_HEADERS,
    N2_DETAIL_ITEM_ID,
    N2_DETAIL_STORAGE_FIELD,
    enrich_n2_detail_row,
    persisted_n2_detail_row,
)
from app.services.n2_taxes_payable_service import _SHEET_CONFIGS, export_data, import_data


class _Result:
    def __init__(self, *, payload: str | None = None, scalar: Any = None) -> None:
        self._row = SimpleNamespace(payload=payload) if payload is not None else None
        self._scalar = scalar

    def fetchone(self) -> Any:
        return self._row

    def scalar_one_or_none(self) -> Any:
        return self._scalar


class _Db:
    def __init__(self, payload: str | None = None) -> None:
        self.payload = payload
        self.calls: list[tuple[str, dict[str, Any]]] = []
        self.flushed = False

    async def execute(self, statement: Any, params: dict[str, Any]) -> _Result:
        sql = str(statement)
        self.calls.append((sql, dict(params)))
        if "SELECT project_id FROM working_paper" in sql:
            return _Result(scalar="00000000-0000-0000-0000-000000000001")
        if "SELECT conclusion AS payload" in sql:
            return _Result(payload=self.payload)
        return _Result()

    async def flush(self) -> None:
        self.flushed = True


def _raw_row() -> dict[str, Any]:
    return {
        "id": "row-1",
        "taxType": "企业所得税",
        "taxRate": 0.25,
        "unadjBegin": 100,
        "unadjPayable": 30,
        "unadjPaid": 10,
        "beginAdjust": 5,
        "beginRje": 7,
        "ajePayable": 2,
        "ajePaid": 3,
        "rjePayable": 4,
        "rjePaid": 1,
        "remark": "核对完成",
    }


def test_n2_detail_formula_uses_both_opening_adjustments_and_overwrites_cache() -> None:
    row = enrich_n2_detail_row({**_raw_row(), "audBegin": 999, "audEnd": 999})
    assert row["unadjEnd"] == 120
    assert row["audBegin"] == 112  # C + G + H
    assert row["audPayable"] == 36
    assert row["audPaid"] == 14
    assert row["audEnd"] == 134


def test_persisted_shape_keeps_begin_rje_but_drops_formula_cache() -> None:
    persisted = persisted_n2_detail_row(
        {**_raw_row(), "unadjEnd": 999, "audBegin": 999, "audEnd": 999},
        identity="fallback",
    )
    assert persisted["id"] == "row-1"
    assert persisted["beginRje"] == 7
    for derived in ("unadjEnd", "audBegin", "audPayable", "audPaid", "audEnd"):
        assert derived not in persisted


def test_n2_detail_contract_is_one_item_and_one_payload_column() -> None:
    from app.routers.wp_render_strategies import _n2_import_export as legacy_factory
    from app.services.workpaper_sync import phase5_n2_taxes_payable as provider
    from app.services.workpaper_sync.store_item_registry import STORE_MERGE_REGISTRY

    config = _SHEET_CONFIGS["N2-2"]
    legacy = legacy_factory._SPECS["N2-2"]
    plan = STORE_MERGE_REGISTRY[provider.ADAPTER_ID]
    assert provider.STORE_ITEM_ID == N2_DETAIL_ITEM_ID
    assert provider.PAYLOAD_COLUMN == N2_DETAIL_STORAGE_FIELD
    assert config["item_id"] == N2_DETAIL_ITEM_ID
    assert config["storage_field"] == N2_DETAIL_STORAGE_FIELD
    assert legacy["item_id"] == N2_DETAIL_ITEM_ID
    assert legacy["storage_field"] == N2_DETAIL_STORAGE_FIELD
    assert tuple(config["headers"]) == N2_DETAIL_HEADERS
    assert tuple(config["fields"]) == N2_DETAIL_FIELDS
    assert plan.item_ids == (N2_DETAIL_ITEM_ID,)


@pytest.mark.asyncio
async def test_export_reads_conclusion_and_recomputes_formula_columns() -> None:
    db = _Db(json.dumps([_raw_row()], ensure_ascii=False))
    buf = await export_data("wp-1", db, sheet="N2-2")
    sql, params = db.calls[0]
    assert "SELECT conclusion AS payload" in sql
    assert params["item_id"] == N2_DETAIL_ITEM_ID

    wb = openpyxl.load_workbook(buf, data_only=True)
    ws = wb["N2-2 明细表"]
    assert [c.value for c in ws[1]] == list(N2_DETAIL_HEADERS)
    values = {field: ws.cell(2, i + 1).value for i, field in enumerate(N2_DETAIL_FIELDS)}
    assert values["unadjEnd"] == 120
    assert values["audBegin"] == 112
    assert values["audEnd"] == 134
    assert values["taxRate"] == 0.25
    wb.close()


@pytest.mark.asyncio
async def test_import_writes_conclusion_with_project_and_only_persisted_fields() -> None:
    wb = openpyxl.Workbook()
    ws = wb.active
    ws.title = "N2-2 明细表"
    ws.append(list(N2_DETAIL_HEADERS))
    full = enrich_n2_detail_row(_raw_row())
    ws.append([full.get(field) for field in N2_DETAIL_FIELDS])
    buf = io.BytesIO()
    wb.save(buf)
    wb.close()

    db = _Db()
    result = await import_data(
        "wp-1",
        buf.getvalue(),
        db,
        sheet="N2-2",
        project_id="00000000-0000-0000-0000-000000000001",
        updated_by="00000000-0000-0000-0000-000000000002",
    )
    assert result["imported_count"] == 1
    assert result["warnings"] == []
    assert result["_item_ids"] == [N2_DETAIL_ITEM_ID]
    assert db.flushed is True
    sql, params = db.calls[-1]
    assert "project_id" in sql
    assert "conclusion" in sql
    assert "updated_by" in sql
    assert "content_version = COALESCE(checklist_responses.content_version, 1) + 1" in sql
    assert params["item_id"] == N2_DETAIL_ITEM_ID
    assert params["updated_by"] == "00000000-0000-0000-0000-000000000002"
    payload = json.loads(params["payload"])
    assert len(payload) == 1 and payload[0]["beginRje"] == 7
    assert payload[0]["taxRate"] == 0.25
    for derived in ("unadjEnd", "audBegin", "audPayable", "audPaid", "audEnd"):
        assert derived not in payload[0]


def _import_workbook(headers: list[str], row: list[Any]) -> bytes:
    wb = openpyxl.Workbook()
    ws = wb.active
    ws.title = "N2-2 明细表"
    ws.append(headers)
    ws.append(row)
    buffer = io.BytesIO()
    wb.save(buffer)
    wb.close()
    return buffer.getvalue()


@pytest.mark.asyncio
async def test_import_rejects_legacy_23_column_header_before_writing() -> None:
    legacy_headers = [
        "税种", "明细子目", "税目", "计税依据", "适用税率", "纳税期间",
        "期初余额", "本期计提", "本期缴纳", "期末余额", "申报表金额", "差异",
        "差异原因", "核查结论", "完税凭证号", "缴纳日期", "纳税地点", "减免标识",
        "减免金额", "计提凭证号", "缴纳凭证号", "关联科目", "备注",
    ]
    db = _Db()

    with pytest.raises(ValueError, match="表头不匹配"):
        await import_data(
            "wp-1",
            _import_workbook(legacy_headers, ["增值税"] * len(legacy_headers)),
            db,
            sheet="N2-2",
            project_id="00000000-0000-0000-0000-000000000001",
            updated_by="00000000-0000-0000-0000-000000000002",
        )

    assert not any("INSERT INTO checklist_responses" in sql for sql, _ in db.calls)


@pytest.mark.asyncio
async def test_import_rejects_non_finite_number_before_writing() -> None:
    row = [""] * len(N2_DETAIL_FIELDS)
    row[0] = "企业所得税"
    row[2] = "1e309"
    db = _Db()

    with pytest.raises(ValueError, match="不是有限数"):
        await import_data(
            "wp-1",
            _import_workbook(list(N2_DETAIL_HEADERS), row),
            db,
            sheet="N2-2",
            project_id="00000000-0000-0000-0000-000000000001",
            updated_by="00000000-0000-0000-0000-000000000002",
        )

    assert not any("INSERT INTO checklist_responses" in sql for sql, _ in db.calls)


@pytest.mark.asyncio
async def test_import_rejects_duplicate_or_reordered_headers() -> None:
    headers = list(N2_DETAIL_HEADERS)
    headers[2] = headers[1]
    db = _Db()

    with pytest.raises(ValueError, match="重复列"):
        await import_data(
            "wp-1",
            _import_workbook(headers, [""] * len(headers)),
            db,
            sheet="N2-2",
            project_id="00000000-0000-0000-0000-000000000001",
            updated_by="00000000-0000-0000-0000-000000000002",
        )


class _Upload:
    filename = "n2.xlsx"

    def __init__(self, content: bytes = b"xlsx") -> None:
        self.content = content
        self.read_count = 0

    async def read(self) -> bytes:
        self.read_count += 1
        return self.content


class _RouteDb:
    def __init__(self, order: list[str] | None = None) -> None:
        self.order = order if order is not None else []
        self.committed = 0
        self.rolled_back = 0

    async def commit(self) -> None:
        self.order.append("commit")
        self.committed += 1

    async def rollback(self) -> None:
        self.order.append("rollback")
        self.rolled_back += 1


@pytest.mark.asyncio
async def test_n2_export_authorizes_project_read_before_service(
    monkeypatch: pytest.MonkeyPatch,
) -> None:
    from app.routers import n2_taxes_payable as router_module

    called = False

    async def deny(*_args, **_kwargs):
        raise HTTPException(status_code=403, detail="无底稿访问权限")

    async def should_not_export(*_args, **_kwargs):
        nonlocal called
        called = True
        return io.BytesIO()

    monkeypatch.setattr(router_module, "authorize_wp_read", deny)
    monkeypatch.setattr(
        "app.services.n2_taxes_payable_service.export_data",
        should_not_export,
    )

    with pytest.raises(HTTPException) as exc:
        await router_module.n2_export_data(
            "00000000-0000-0000-0000-000000000010",
            sheet="N2-2",
            db=_RouteDb(),
            current_user=SimpleNamespace(id=UUID(int=3)),
        )

    assert exc.value.status_code == 403
    assert called is False


@pytest.mark.asyncio
async def test_n2_import_authorizes_and_publishes_only_after_commit(
    monkeypatch: pytest.MonkeyPatch,
) -> None:
    from app.routers import checklist_responses, n2_taxes_payable as router_module

    wp_id = UUID("00000000-0000-0000-0000-000000000010")
    project_id = UUID("00000000-0000-0000-0000-000000000020")
    user_id = UUID("00000000-0000-0000-0000-000000000030")
    order: list[str] = []
    gate_call: dict[str, Any] = {}
    service_call: dict[str, Any] = {}

    async def authorize(*_args, **_kwargs):
        order.append("authorize")
        return project_id

    async def gate(*_args, **kwargs):
        order.append("gate")
        gate_call.update(kwargs)

    async def import_rows(*args, **kwargs):
        order.append("write")
        service_call.update(kwargs)
        return {
            "imported_count": 1,
            "imported_counts": {"N2-2": 1},
            "warnings": [],
            "message": "导入成功",
            "_item_ids": [N2_DETAIL_ITEM_ID],
        }

    async def publish(**kwargs):
        order.append("publish")
        assert kwargs["item_ids"] == [N2_DETAIL_ITEM_ID]
        assert kwargs["user_id"] == user_id

    monkeypatch.setattr(router_module, "authorize_wp_edit", authorize)
    monkeypatch.setattr(router_module, "enforce_wp_gate", gate)
    # _n2_wp_code 是 async，使用协程替身。
    async def wp_code(*_args):
        return "N2"

    async def audit_year(*_args):
        return 2025

    monkeypatch.setattr(router_module, "_n2_wp_code", wp_code)
    monkeypatch.setattr(router_module, "fetch_project_audit_year", audit_year)
    monkeypatch.setattr(
        "app.services.n2_taxes_payable_service.import_data",
        import_rows,
    )
    monkeypatch.setattr(checklist_responses, "_publish_checklist_saved", publish)

    db = _RouteDb(order)
    upload = _Upload()
    result = await router_module.n2_import_data(
        str(wp_id),
        sheet="N2-2",
        file=upload,
        db=db,
        current_user=SimpleNamespace(id=user_id),
    )

    assert result == {
        "imported_count": 1,
        "imported_counts": {"N2-2": 1},
        "warnings": [],
        "message": "导入成功",
    }
    assert order == ["authorize", "gate", "write", "commit", "publish"]
    assert upload.read_count == 1
    assert gate_call["action"] == "save_checklist"
    assert gate_call["project_id"] == project_id
    assert gate_call.get("requested_sheet_key") is None
    assert service_call["project_id"] == project_id
    assert service_call["updated_by"] == user_id
    assert db.rolled_back == 0


@pytest.mark.asyncio
async def test_n2_import_denial_happens_before_upload_read(
    monkeypatch: pytest.MonkeyPatch,
) -> None:
    from app.routers import n2_taxes_payable as router_module

    async def deny(*_args, **_kwargs):
        raise HTTPException(status_code=403, detail="无底稿编辑权限")

    monkeypatch.setattr(router_module, "authorize_wp_edit", deny)
    upload = _Upload()
    with pytest.raises(HTTPException) as exc:
        await router_module.n2_import_data(
            "00000000-0000-0000-0000-000000000010",
            sheet="N2-2",
            file=upload,
            db=_RouteDb(),
            current_user=SimpleNamespace(id=UUID(int=3)),
        )

    assert exc.value.status_code == 403
    assert upload.read_count == 0

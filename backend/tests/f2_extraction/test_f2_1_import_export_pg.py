"""F2-1 存货审定表导入 / 导出 —— 真实 PostgreSQL + 真实应用路由（含 router 级 dedicated_wp_gate）。

两条路径原先都在函数体内 import 全仓不存在的 ``app.models.checklist_response_models``，一调即
ModuleNotFoundError（lazy import 守卫遇「整个模块缺失」直接跳过，所以一直没打红）。``checklist_responses``
没有 ORM 模型，两条路径改走裸 SQL；``wp_id`` 是 uuid 列、写入靠 ``(wp_id, item_id)`` 唯一约束上的
ON CONFLICT，这两点只有 PG 能证明，故本文件只在 PG 上跑。

每个用例在一个最终回滚的事务里建数据；会话用 ``create_savepoint`` 吸收路由里的 commit，不污染开发库。
"""

from __future__ import annotations

import io

import pytest
import pytest_asyncio
import sqlalchemy as sa
from openpyxl import Workbook, load_workbook
from sqlalchemy.ext.asyncio import AsyncSession, create_async_engine

from app.core.config import settings
from app.routers.wp_render_strategies._f2_import_export import _F2_1_HEADERS
from tests.procedure_delegation_visibility._factories import mk_project, mk_user, mk_working_paper, mk_wp_index

IS_PG = settings.DATABASE_URL.startswith("postgresql")
pytestmark = [pytest.mark.pg_only, pytest.mark.skipif(not IS_PG, reason="need PostgreSQL")]

_XLSX = "application/vnd.openxmlformats-officedocument.spreadsheetml.sheet"


@pytest_asyncio.fixture
async def pg_session():
    engine = create_async_engine(settings.DATABASE_URL, pool_pre_ping=True)
    conn = await engine.connect()
    trans = await conn.begin()
    session = AsyncSession(bind=conn, expire_on_commit=False, join_transaction_mode="create_savepoint")
    try:
        yield session
    finally:
        await session.close()
        await trans.rollback()
        await conn.close()
        await engine.dispose()


async def _f2_paper(session):
    project = await mk_project(session)
    index = await mk_wp_index(session, project.id, wp_code="F2", audit_cycle="F")
    return project, await mk_working_paper(session, project.id, index.id)


def _workbook(gross: list[list], impairment: list[list]) -> bytes:
    wb = Workbook()
    wb.active.title = "编制说明"
    for title, rows in (("原值", gross), ("跌价准备", impairment)):
        ws = wb.create_sheet(title)
        ws.append(_F2_1_HEADERS)
        for row in rows:
            ws.append(row)
    buffer = io.BytesIO()
    wb.save(buffer)
    return buffer.getvalue()


async def _stored(session, wp_id) -> dict:
    result = await session.execute(
        sa.text("SELECT item_id, conclusion, project_id FROM checklist_responses WHERE wp_id = :w"),
        {"w": wp_id},
    )
    return {row.item_id: row for row in result.fetchall()}


def _body(resp) -> dict:
    payload = resp.json()
    return payload.get("data", payload)


async def test_f2_1_import_writes_one_row_per_field(pg_session):
    from app.main import app
    from tests._test_auth_helper import FakeAuthUser, override_auth

    project, paper = await _f2_paper(pg_session)
    admin = await mk_user(pg_session, role="admin")
    url = f"/api/workpapers/{paper.id}/f2/import-data"
    first = _workbook(
        [["原材料", "1401", 100.5, 50, 30, 10], ["不存在的类别", "", 1, 1, 1, 1]],
        [["存货跌价准备", "1471", 20, 5, 3, 2]],
    )
    # 再导入同一行：走真实唯一约束上的 ON CONFLICT；空值不覆盖已有数据（模板编制说明第三条）
    second = _workbook([["原材料", "1401", 200, None, None, None]], [])
    async with override_auth(app, db_session=pg_session, user=FakeAuthUser(user_id=admin.id)) as client:
        r1 = await client.post(url, params={"sheet": "F2-1"}, files={"file": ("f2.xlsx", first, _XLSX)})
        assert r1.status_code == 200, r1.text
        r2 = await client.post(url, params={"sheet": "F2-1"}, files={"file": ("f2.xlsx", second, _XLSX)})
        assert r2.status_code == 200, r2.text

    b1, b2 = _body(r1), _body(r2)
    assert (b1["imported_count"], b1["field_count"], b1["skipped_count"]) == (2, 8, 1)
    assert "未知类别「不存在的类别」已跳过" in b1["warnings"]
    assert (b2["imported_count"], b2["field_count"]) == (1, 1)
    stored = await _stored(pg_session, paper.id)
    assert len(stored) == 8 and {row.project_id for row in stored.values()} == {project.id}
    assert stored["F2-1-gross-raw-materials-opening"].conclusion == "200.0"
    assert stored["F2-1-gross-raw-materials-increase"].conclusion == "50.0"
    assert stored["F2-1-impairment-impairment-provision-adjustment"].conclusion == "2.0"


async def test_f2_1_export_reads_only_this_papers_f2_1_rows(pg_session):
    from app.main import app
    from tests._test_auth_helper import FakeAuthUser, override_auth

    project, paper = await _f2_paper(pg_session)
    other_project, other = await _f2_paper(pg_session)
    admin = await mk_user(pg_session, role="admin")
    seed = [
        (project.id, paper.id, "F2-1-gross-raw-materials-opening", "200"),
        (project.id, paper.id, "F2-1-gross-raw-materials-decrease", "30.5"),
        (project.id, paper.id, "F2-1-impairment-impairment-provision-increase", "5"),
        (project.id, paper.id, "F2-3-rows", "999"),  # 别的 sheet：LIKE 'F2-1-%' 不命中
        (other_project.id, other.id, "F2-1-gross-raw-materials-increase", "777"),  # 别的底稿
    ]
    for project_id, wp_id, item_id, value in seed:
        await pg_session.execute(
            sa.text("INSERT INTO checklist_responses (project_id, wp_id, item_id, conclusion) VALUES (:p, :w, :i, :v)"),
            {"p": project_id, "w": wp_id, "i": item_id, "v": value},
        )
    async with override_auth(app, db_session=pg_session, user=FakeAuthUser(user_id=admin.id)) as client:
        resp = await client.post(f"/api/workpapers/{paper.id}/f2/export-data", params={"sheet": "F2-1"})
    assert resp.status_code == 200, resp.text

    wb = load_workbook(io.BytesIO(resp.content), data_only=True)
    gross = {row[0]: row[2:6] for row in wb["原值"].iter_rows(min_row=2, values_only=True)}
    assert gross["原材料"] == (200.0, None, 30.5, None)
    assert all(values == (None, None, None, None) for label, values in gross.items() if label != "原材料")
    impairment = {row[0]: row[2:6] for row in wb["跌价准备"].iter_rows(min_row=2, values_only=True)}
    assert impairment == {"存货跌价准备": (None, 5.0, None, None)}

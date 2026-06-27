"""Property-Based Tests for word-template render strategy.

Property 5: API response merges checklist_responses correctly
Property 12: Render strategy response structure

**Validates: Requirements 3.4, 4.4, 10.1, 10.2, 10.4**
"""

from __future__ import annotations

import asyncio
import os
import tempfile
from dataclasses import dataclass, field
from unittest.mock import AsyncMock, MagicMock

import pytest
from docx import Document
from hypothesis import HealthCheck, given, settings
from hypothesis import strategies as st

from app.routers.wp_render_strategies._context import RenderContext
from app.routers.wp_render_strategies._word_template import render


# ─── Fake dataclasses (mirror parser output) ──────────────────────────────────


@dataclass
class FakePlaceholder:
    field_id: str
    label: str
    data_type: str
    default_value: str
    position: dict
    pattern: str


@dataclass
class FakeParagraph:
    index: int
    text: str
    style: str
    heading_level: int
    placeholder_ids: list = field(default_factory=list)


@dataclass
class FakeTable:
    index: int
    rows: list = field(default_factory=list)
    placeholder_ids: list = field(default_factory=list)


@dataclass
class FakeStructure:
    paragraphs: list = field(default_factory=list)
    tables: list = field(default_factory=list)
    placeholders: list = field(default_factory=list)
    metadata: dict = field(default_factory=dict)


# ─── Hypothesis Strategies ────────────────────────────────────────────────────

_field_id_st = st.from_regex(r"[a-z][a-z0-9_]{2,15}", fullmatch=True)
_wp_code_st = st.from_regex(r"A[0-9]{1,2}-[0-9]", fullmatch=True)
_value_st = st.text(min_size=1, max_size=30, alphabet=st.characters(categories=("L", "N", "P")))


@st.composite
def responses_strategy(draw: st.DrawFn) -> tuple[str, list[tuple[str, str]]]:
    """Generate a (wp_code, [(field_id, value), ...]) pair."""
    wp_code = draw(_wp_code_st)
    count = draw(st.integers(min_value=1, max_value=5))
    field_ids = draw(st.lists(_field_id_st, min_size=count, max_size=count, unique=True))
    values = draw(st.lists(_value_st, min_size=count, max_size=count))
    return wp_code, list(zip(field_ids, values))


@st.composite
def structure_with_placeholders_strategy(draw: st.DrawFn) -> tuple[str, FakeStructure, list[tuple[str, str]]]:
    """Generate (wp_code, FakeStructure, responses) for render testing."""
    wp_code = draw(_wp_code_st)
    count = draw(st.integers(min_value=1, max_value=5))
    field_ids = draw(st.lists(_field_id_st, min_size=count, max_size=count, unique=True))
    values = draw(st.lists(_value_st, min_size=count, max_size=count))
    responses = list(zip(field_ids, values))

    placeholders = [
        FakePlaceholder(
            field_id=fid,
            label=f"标签_{fid}",
            data_type="text",
            default_value=f"默认_{fid}",
            position={"paragraph_index": i},
            pattern=f"${{{fid}}}",
        )
        for i, fid in enumerate(field_ids)
    ]
    paragraphs = [
        FakeParagraph(
            index=i,
            text=f"段落含 ${{{fid}}} 内容",
            style="Normal",
            heading_level=0,
            placeholder_ids=[fid],
        )
        for i, fid in enumerate(field_ids)
    ]
    structure = FakeStructure(
        paragraphs=paragraphs,
        tables=[],
        placeholders=placeholders,
        metadata={"template_name": f"{wp_code} 模板", "wp_code": wp_code, "last_parsed_at": "2026-01-01T00:00:00Z"},
    )
    return wp_code, structure, responses


# ─── Helper: build mock RenderContext ─────────────────────────────────────────


def _make_ctx(
    wp_code: str,
    structure: FakeStructure,
    responses: list[tuple[str, str]],
    template_file_path: str = "/fake/template.docx",
) -> RenderContext:
    """Build a mock RenderContext with given structure and checklist responses."""
    prefix = f"wt-{wp_code}-"

    # Build fake DB rows
    rows = []
    for field_id, value in responses:
        row = MagicMock()
        row.item_id = f"{prefix}{field_id}"
        row.conclusion = value
        row.remark = ""
        rows.append(row)

    db = AsyncMock()
    result_mock = MagicMock()
    result_mock.fetchall.return_value = rows
    db.execute = AsyncMock(return_value=result_mock)

    ctx = MagicMock(spec=RenderContext)
    ctx.wp_id = "wp-test-001"
    ctx.wp_code = wp_code
    ctx.db = db
    ctx.template_file_path = template_file_path
    return ctx


# ─── Property 5: Response merge ──────────────────────────────────────────────


class TestProperty5ResponseMerge:
    """Property 5: API response merges checklist_responses correctly.

    For any set of checklist_responses stored with item_id pattern
    `wt-{wp_code}-{field_id}`, the render response SHALL include each
    stored value as `current_value` on the corresponding placeholder.

    **Validates: Requirements 3.4, 4.4**
    """

    @settings(max_examples=5, suppress_health_check=[HealthCheck.function_scoped_fixture])
    @given(data=structure_with_placeholders_strategy())
    def test_all_responses_merged_as_current_value(self, data):
        """Every stored checklist_response value appears as current_value on its placeholder."""
        wp_code, structure, responses = data
        ctx = _make_ctx(wp_code, structure, responses)

        from unittest.mock import patch

        with patch(
            "app.services.wp_docx_template_parser.get_cached_structure",
            return_value=structure,
        ):
            result = asyncio.get_event_loop().run_until_complete(render(ctx))

        assert result is not None
        ts = result["template_structure"]
        placeholders = ts["placeholders"]

        # Build a lookup: field_id → current_value from response
        response_map = dict(responses)
        placeholder_map = {p["field_id"]: p["current_value"] for p in placeholders}

        for field_id, expected_value in response_map.items():
            assert field_id in placeholder_map, (
                f"field_id '{field_id}' not found in template_structure placeholders"
            )
            assert placeholder_map[field_id] == expected_value, (
                f"current_value mismatch for '{field_id}': "
                f"expected '{expected_value}', got '{placeholder_map[field_id]}'"
            )

    @settings(max_examples=5, suppress_health_check=[HealthCheck.function_scoped_fixture])
    @given(data=structure_with_placeholders_strategy())
    def test_filled_responses_dict_matches_stored_values(self, data):
        """filled_responses dict contains exactly the stored field_id→value pairs."""
        wp_code, structure, responses = data
        ctx = _make_ctx(wp_code, structure, responses)

        from unittest.mock import patch

        with patch(
            "app.services.wp_docx_template_parser.get_cached_structure",
            return_value=structure,
        ):
            result = asyncio.get_event_loop().run_until_complete(render(ctx))

        assert result is not None
        filled = result["filled_responses"]

        for field_id, expected_value in responses:
            assert field_id in filled, f"field_id '{field_id}' not in filled_responses"
            assert filled[field_id] == expected_value


# ─── Property 12: Render strategy response structure ─────────────────────────


class TestProperty12RenderStrategyResponseStructure:
    """Property 12: Render strategy response structure.

    For any valid word-template render, response SHALL contain
    `template_structure` (with placeholders/paragraphs/tables/metadata),
    `filled_responses` (dict), and `sign_status`.

    **Validates: Requirements 10.1, 10.2, 10.4**
    """

    @settings(max_examples=5, suppress_health_check=[HealthCheck.function_scoped_fixture])
    @given(data=structure_with_placeholders_strategy())
    def test_response_has_three_top_level_keys(self, data):
        """Response dict contains template_structure, filled_responses, sign_status."""
        wp_code, structure, responses = data
        ctx = _make_ctx(wp_code, structure, responses)

        from unittest.mock import patch

        with patch(
            "app.services.wp_docx_template_parser.get_cached_structure",
            return_value=structure,
        ):
            result = asyncio.get_event_loop().run_until_complete(render(ctx))

        assert result is not None
        expected_keys = {"template_structure", "filled_responses", "sign_status"}
        assert set(result.keys()) == expected_keys, (
            f"Top-level keys mismatch: {set(result.keys())} != {expected_keys}"
        )

    @settings(max_examples=5, suppress_health_check=[HealthCheck.function_scoped_fixture])
    @given(data=structure_with_placeholders_strategy())
    def test_template_structure_has_required_subkeys(self, data):
        """template_structure contains placeholders, paragraphs, tables, metadata."""
        wp_code, structure, responses = data
        ctx = _make_ctx(wp_code, structure, responses)

        from unittest.mock import patch

        with patch(
            "app.services.wp_docx_template_parser.get_cached_structure",
            return_value=structure,
        ):
            result = asyncio.get_event_loop().run_until_complete(render(ctx))

        assert result is not None
        ts = result["template_structure"]
        assert isinstance(ts, dict)
        assert "placeholders" in ts
        assert "paragraphs" in ts
        assert "tables" in ts
        assert "metadata" in ts
        assert isinstance(ts["placeholders"], list)
        assert isinstance(ts["paragraphs"], list)
        assert isinstance(ts["tables"], list)
        assert isinstance(ts["metadata"], dict)

    @settings(max_examples=5, suppress_health_check=[HealthCheck.function_scoped_fixture])
    @given(data=structure_with_placeholders_strategy())
    def test_filled_responses_is_dict_and_sign_status_exists(self, data):
        """filled_responses is a dict of field_id→value, sign_status key exists."""
        wp_code, structure, responses = data
        ctx = _make_ctx(wp_code, structure, responses)

        from unittest.mock import patch

        with patch(
            "app.services.wp_docx_template_parser.get_cached_structure",
            return_value=structure,
        ):
            result = asyncio.get_event_loop().run_until_complete(render(ctx))

        assert result is not None
        assert isinstance(result["filled_responses"], dict)
        # sign_status can be None (placeholder for future) or dict
        assert "sign_status" in result

    @settings(max_examples=5, suppress_health_check=[HealthCheck.function_scoped_fixture])
    @given(data=structure_with_placeholders_strategy())
    def test_each_placeholder_has_current_value_key(self, data):
        """Every placeholder in template_structure has a 'current_value' key."""
        wp_code, structure, responses = data
        ctx = _make_ctx(wp_code, structure, responses)

        from unittest.mock import patch

        with patch(
            "app.services.wp_docx_template_parser.get_cached_structure",
            return_value=structure,
        ):
            result = asyncio.get_event_loop().run_until_complete(render(ctx))

        assert result is not None
        for p in result["template_structure"]["placeholders"]:
            assert "current_value" in p, (
                f"Placeholder '{p.get('field_id', '?')}' missing 'current_value' key"
            )

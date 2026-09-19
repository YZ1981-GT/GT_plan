"""F2-23 监盘小结 docx ↔ 结构化字段同步."""

from __future__ import annotations

from pathlib import Path

from app.services.f2_stocktake_summary_sync import (
    create_g2_6_1_template_docx,
    extract_fields_from_docx,
    fill_summary_docx,
    merge_extracted_into_existing,
    migrate_legacy_fields,
)


def test_f2_23_template_has_placeholders(tmp_path: Path) -> None:
    tpl = create_g2_6_1_template_docx(tmp_path / "t.docx")
    from docx import Document

    text = "\n".join(p.text for p in Document(str(tpl)).paragraphs)
    assert "${purpose}" in text
    assert "${resultByLocation}" in text
    assert "${summaryBlock}" in text
    assert "G2-6-1" in text


def test_f2_23_fill_extract_roundtrip(tmp_path: Path) -> None:
    tpl = create_g2_6_1_template_docx(tmp_path / "tpl.docx")
    target = tmp_path / "filled.docx"
    sample = {
        "entityName": "测试股份有限公司",
        "auditYear": "2025",
        "bsDate": "2025-12-31",
        "purpose": "确定存货存在并为公司所有。",
        "scope": "公司本部存货；少量分厂拟审计时盘点。",
        "warehouses": "一分厂、工业园、造纸厂",
        "countDate": "2025年12月31日至2026年1月1日",
        "clientStaff": "仓管、物资管理、财务监督",
        "auditors": "张三、李四",
        "assignment": "张三：一分厂；李四：工业园",
        "countMethod": "主要原料采用实地盘存制。",
        "processOverview": "对一分厂、工业园全部产品及原料监盘。",
        "inventoryTotal": "1200",
        "sampleAmount": "900",
        "samplePct": "75",
        "coverageNote": "因公司期末未结账，金额为暂估",
        "resultByLocation": "1).一分厂\n抽盘结果无差异。详见一分厂存货抽盘表。",
        "overallConclusion": "存货管理较好，未见重大账实不符。",
        "followUp": "关注造纸厂散装计量差异。",
        "summaryDate": "2026-01-02",
        "teamName": "测试审计小组",
    }
    fill_summary_docx(tpl, target, sample)
    extracted = extract_fields_from_docx(target)
    for key in (
        "entityName",
        "auditYear",
        "purpose",
        "scope",
        "warehouses",
        "countMethod",
        "resultByLocation",
        "overallConclusion",
        "teamName",
        "summaryDate",
    ):
        assert extracted.get(key) == sample[key], f"{key}: {extracted.get(key)!r} != {sample[key]!r}"
    assert extracted.get("bsDate") == "2025-12-31"
    assert "1200" in (extracted.get("inventoryTotal") or extracted.get("processOverview") or "")


def test_migrate_legacy_and_merge() -> None:
    legacy = migrate_legacy_fields(
        {"inventoryComposition": "实地盘存", "purpose": "旧目的", "customExtra": "保留"}
    )
    assert legacy["countMethod"] == "实地盘存"
    merged = merge_extracted_into_existing(legacy, {"purpose": "新目的", "scope": "新范围"})
    assert merged["purpose"] == "新目的"
    assert merged["scope"] == "新范围"
    assert merged["customExtra"] == "保留"
    assert merged["countMethod"] == "实地盘存"

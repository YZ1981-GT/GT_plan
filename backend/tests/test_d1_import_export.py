"""Property 13 (backend): D1 导入导出 Round-Trip + 模板格式校验

Feature: d1-adjudication-table, Property 13: 动态行序列化Round-Trip

Tests:
1. Round-trip: generate random row data → export → import → verify equivalent
2. Template format validation: random column arrangement → verify error detection

Uses hypothesis with max_examples=5 (per project convention for PBT).

**Validates: Requirements 9.3, 9.4**
"""

from __future__ import annotations

import sys
from pathlib import Path

from hypothesis import given, settings
from hypothesis import strategies as st

# Ensure the backend app is importable
sys.path.insert(0, str(Path(__file__).resolve().parent.parent))

from app.routers.wp_render_strategies._d1_import_export import (
    _export_d1_2_row,
    _export_d1_3_row,
    _export_d1_4_row,
    _parse_d1_2_row,
    _parse_d1_3_row,
    _parse_d1_4_row,
    _validate_columns,
    _create_template_wb,
    _SHEET_HEADERS,
)

# ═══════════════════════════════════════════════════════════════════════════════
# Strategies
# ═══════════════════════════════════════════════════════════════════════════════

_amount_st = st.floats(min_value=-1e9, max_value=1e9, allow_nan=False, allow_infinity=False)
_text_st = st.text(min_size=1, max_size=20, alphabet=st.characters(categories=("L", "N")))
_category_d1_2_st = st.sampled_from(["银行承兑汇票", "商业承兑汇票", "信用证", "商业汇票"])
_relation_st = st.sampled_from(["关联方", "非关联方", ""])

_d1_2_row_st = st.fixed_dictionaries({
    "rowId": st.uuids().map(str),
    "category": _category_d1_2_st,
    "isFixed": st.booleans(),
    "priorUnadjusted": _amount_st,
    "priorAje": _amount_st,
    "priorRje": _amount_st,
    "currentIncrease": _amount_st,
    "currentDecrease": _amount_st,
    "currentAje": _amount_st,
    "currentRje": _amount_st,
})

_d1_3_row_st = st.fixed_dictionaries({
    "rowId": st.uuids().map(str),
    "customerName": _text_st,
    "companyCode": st.text(min_size=0, max_size=10, alphabet=st.characters(categories=("L", "N"))),
    "relationType": _relation_st,
    "priorUnadjusted": _amount_st,
    "priorAje": _amount_st,
    "priorRje": _amount_st,
    "currentIncrease": _amount_st,
    "currentDecrease": _amount_st,
    "currentBalance": _amount_st,
    "reclassification": _amount_st,
    "currentAje": _amount_st,
    "currentRje": _amount_st,
})

_d1_4_row_st = st.fixed_dictionaries({
    "rowId": st.uuids().map(str),
    "label": _text_st,
    "category": st.sampled_from(["individual", "portfolio"]),
    "isSubRow": st.just(True),
    "priorUnadjusted": _amount_st,
    "priorAje": _amount_st,
    "priorRje": _amount_st,
    "currentProvision": _amount_st,
    "currentRecovery": _amount_st,
    "currentReversal": _amount_st,
    "currentWriteOff": _amount_st,
    "currentOther": _amount_st,
    "currentAje": _amount_st,
    "currentRje": _amount_st,
})


# ═══════════════════════════════════════════════════════════════════════════════
# Round-Trip Tests
# ═══════════════════════════════════════════════════════════════════════════════


@settings(max_examples=5)
@given(rows=st.lists(_d1_2_row_st, min_size=1, max_size=5))
def test_d1_2_export_import_round_trip(rows: list[dict]) -> None:
    """D1-2 按类别明细表 export → import round-trip preserves key fields.

    **Validates: Requirements 9.3**
    """
    headers = _SHEET_HEADERS["D1-2"]

    for original in rows:
        exported = _export_d1_2_row(original)
        assert len(exported) == len(headers)

        imported = _parse_d1_2_row(tuple(exported), headers)

        # String fields
        assert imported["category"] == original["category"]

        # Numeric fields
        assert abs(imported["priorUnadjusted"] - float(original["priorUnadjusted"])) < 1e-6
        assert abs(imported["priorAje"] - float(original["priorAje"])) < 1e-6
        assert abs(imported["priorRje"] - float(original["priorRje"])) < 1e-6
        assert abs(imported["currentIncrease"] - float(original["currentIncrease"])) < 1e-6
        assert abs(imported["currentDecrease"] - float(original["currentDecrease"])) < 1e-6
        assert abs(imported["currentAje"] - float(original["currentAje"])) < 1e-6
        assert abs(imported["currentRje"] - float(original["currentRje"])) < 1e-6

        # rowId is regenerated (new UUID)
        assert imported["rowId"] != ""


@settings(max_examples=5)
@given(rows=st.lists(_d1_3_row_st, min_size=1, max_size=5))
def test_d1_3_export_import_round_trip(rows: list[dict]) -> None:
    """D1-3 按客户明细表 export → import round-trip preserves key fields.

    **Validates: Requirements 9.3**
    """
    headers = _SHEET_HEADERS["D1-3"]

    for original in rows:
        exported = _export_d1_3_row(original)
        assert len(exported) == len(headers)

        imported = _parse_d1_3_row(tuple(exported), headers)

        # String fields
        assert imported["customerName"] == original["customerName"]
        assert imported["companyCode"] == original["companyCode"]
        assert imported["relationType"] == original["relationType"]

        # Numeric fields
        assert abs(imported["priorUnadjusted"] - float(original["priorUnadjusted"])) < 1e-6
        assert abs(imported["priorAje"] - float(original["priorAje"])) < 1e-6
        assert abs(imported["priorRje"] - float(original["priorRje"])) < 1e-6
        assert abs(imported["currentIncrease"] - float(original["currentIncrease"])) < 1e-6
        assert abs(imported["currentDecrease"] - float(original["currentDecrease"])) < 1e-6
        assert abs(imported["currentBalance"] - float(original["currentBalance"])) < 1e-6
        assert abs(imported["reclassification"] - float(original["reclassification"])) < 1e-6
        assert abs(imported["currentAje"] - float(original["currentAje"])) < 1e-6
        assert abs(imported["currentRje"] - float(original["currentRje"])) < 1e-6

        # rowId is regenerated (new UUID)
        assert imported["rowId"] != ""


@settings(max_examples=5)
@given(rows=st.lists(_d1_4_row_st, min_size=1, max_size=5))
def test_d1_4_export_import_round_trip(rows: list[dict]) -> None:
    """D1-4 坏账准备明细表 export → import round-trip preserves key fields.

    **Validates: Requirements 9.3**
    """
    headers = _SHEET_HEADERS["D1-4"]

    for original in rows:
        exported = _export_d1_4_row(original)
        assert len(exported) == len(headers)

        imported = _parse_d1_4_row(tuple(exported), headers)

        # String fields
        assert imported["label"] == original["label"]

        # Numeric fields
        assert abs(imported["priorUnadjusted"] - float(original["priorUnadjusted"])) < 1e-6
        assert abs(imported["priorAje"] - float(original["priorAje"])) < 1e-6
        assert abs(imported["priorRje"] - float(original["priorRje"])) < 1e-6
        assert abs(imported["currentProvision"] - float(original["currentProvision"])) < 1e-6
        assert abs(imported["currentRecovery"] - float(original["currentRecovery"])) < 1e-6
        assert abs(imported["currentReversal"] - float(original["currentReversal"])) < 1e-6
        assert abs(imported["currentWriteOff"] - float(original["currentWriteOff"])) < 1e-6
        assert abs(imported["currentOther"] - float(original["currentOther"])) < 1e-6
        assert abs(imported["currentAje"] - float(original["currentAje"])) < 1e-6
        assert abs(imported["currentRje"] - float(original["currentRje"])) < 1e-6

        # rowId is regenerated (new UUID)
        assert imported["rowId"] != ""


# ═══════════════════════════════════════════════════════════════════════════════
# Template Format Validation Tests
# ═══════════════════════════════════════════════════════════════════════════════


@settings(max_examples=5)
@given(
    sheet=st.sampled_from(["D1-2", "D1-3", "D1-4"]),
    extra_cols=st.lists(
        st.text(min_size=1, max_size=10, alphabet=st.characters(categories=("L",))),
        min_size=1,
        max_size=3,
    ),
)
def test_validate_columns_detects_missing(sheet: str, extra_cols: list[str]) -> None:
    """Template format validation: random column removal → detect missing columns.

    **Validates: Requirements 9.4**
    """
    expected_headers = _SHEET_HEADERS[sheet]

    # Create a worksheet with only a subset of columns (remove last N and add random)
    from openpyxl import Workbook

    wb = Workbook()
    ws = wb.active

    # Remove some expected columns and replace with random
    partial_headers = expected_headers[:-2] + extra_cols
    for col_idx, col_name in enumerate(partial_headers, 1):
        ws.cell(row=1, column=col_idx, value=col_name)

    missing = _validate_columns(ws, sheet)

    # The removed columns should be detected as missing
    removed = set(expected_headers) - set(partial_headers)
    for col in removed:
        assert col in missing, f"Expected '{col}' to be detected as missing"


def test_template_workbook_has_correct_headers() -> None:
    """Verify template workbook has all expected column headers.

    **Validates: Requirements 9.1**
    """
    for sheet_code in ["D1-2", "D1-3", "D1-4"]:
        wb = _create_template_wb(sheet_code)
        ws = wb.active

        actual_headers = [
            cell.value for cell in next(ws.iter_rows(min_row=1, max_row=1))
            if cell.value is not None
        ]

        expected = _SHEET_HEADERS[sheet_code]
        assert actual_headers == expected, f"{sheet_code}: headers mismatch"
        assert ws.freeze_panes == "A2", f"{sheet_code}: freeze panes not set"


# ═══════════════════════════════════════════════════════════════════════════════
# D1-6/D1-7/D1-8/D1-9 背书贴现组 — Round-Trip + 模板格式 (Task 14.2)
# ═══════════════════════════════════════════════════════════════════════════════

from app.routers.wp_render_strategies._d1_import_export import (  # noqa: E402
    _export_d1_9_row,
    _parse_d1_9_row,
    _D1_7_HEADERS,
    _D1_9_HEADERS,
)

_date_st = st.dates(min_value=__import__("datetime").date(2020, 1, 1),
                    max_value=__import__("datetime").date(2030, 12, 31)).map(
    lambda d: d.isoformat()
)
_note_type_d1_9_st = st.sampled_from(["银行承兑汇票", "商业承兑汇票"])
_rate_st = st.floats(min_value=0.0, max_value=0.5, allow_nan=False, allow_infinity=False)

_d1_9_row_st = st.fixed_dictionaries({
    "rowId": st.uuids().map(str),
    "rowType": st.just("dynamic"),
    "noteType": _note_type_d1_9_st,
    "faceValue": st.floats(min_value=0.0, max_value=1e9, allow_nan=False, allow_infinity=False),
    "faceRate": _rate_st,
    "issueDate": _date_st,
    "maturityDate": _date_st,
    "maturityValue": st.floats(min_value=0.0, max_value=1e9, allow_nan=False, allow_infinity=False),
    "discountDate": _date_st,
    "discountRate": _rate_st,
    "bookedInterest": _amount_st,
    "remark": st.text(min_size=0, max_size=20, alphabet=st.characters(categories=("L", "N"))),
})


@settings(max_examples=5)
@given(rows=st.lists(_d1_9_row_st, min_size=1, max_size=5))
def test_d1_9_export_import_round_trip(rows: list[dict]) -> None:
    """D1-9 贴息检查表 export → import round-trip preserves editable fields.

    派生列（贴息天数/应计贴现利息/差异）导出时计算供参考，导入时忽略。

    **Validates: Requirements 14.3, 14.4**
    """
    headers = _D1_9_HEADERS

    for original in rows:
        exported = _export_d1_9_row(original)
        assert len(exported) == len(headers)

        imported = _parse_d1_9_row(tuple(exported), headers)

        # Editable string fields survive round-trip
        assert imported["noteType"] == original["noteType"]
        assert imported["issueDate"] == original["issueDate"]
        assert imported["maturityDate"] == original["maturityDate"]
        assert imported["discountDate"] == original["discountDate"]
        assert imported["remark"] == original["remark"]

        # Editable numeric fields survive round-trip
        assert abs(imported["faceValue"] - float(original["faceValue"])) < 1e-6
        assert abs(imported["faceRate"] - float(original["faceRate"])) < 1e-6
        assert abs(imported["maturityValue"] - float(original["maturityValue"])) < 1e-6
        assert abs(imported["discountRate"] - float(original["discountRate"])) < 1e-6
        assert abs(imported["bookedInterest"] - float(original["bookedInterest"])) < 1e-6

        # Derived columns are NOT persisted on import
        assert "discountDays" not in imported
        assert "calculatedInterest" not in imported
        assert "difference" not in imported

        # rowId is regenerated
        assert imported["rowId"] != ""


def test_d1_7_template_has_31_columns() -> None:
    """D1-7 备查簿模板必须有恰好31列表头。

    **Validates: Requirements 14.1**
    """
    wb = _create_template_wb("D1-7")
    ws = wb.active

    actual_headers = [
        cell.value for cell in next(ws.iter_rows(min_row=1, max_row=1))
        if cell.value is not None
    ]

    assert len(actual_headers) == 31, f"D1-7应有31列，实际{len(actual_headers)}列"
    assert actual_headers == _D1_7_HEADERS
    assert ws.freeze_panes == "A2"


def test_d1_9_template_has_13_columns_no_data() -> None:
    """D1-9 贴息模板必须有13列表头且无数据行（仅表头）。

    **Validates: Requirements 14.1**
    """
    wb = _create_template_wb("D1-9")
    ws = wb.active

    all_rows = list(ws.iter_rows(values_only=True))
    # Only the header row present
    assert len(all_rows) == 1, f"D1-9模板应只有表头行，实际{len(all_rows)}行"

    header_row = [c for c in all_rows[0] if c is not None]
    assert len(header_row) == 13, f"D1-9应有13列，实际{len(header_row)}列"
    assert header_row == _D1_9_HEADERS


@settings(max_examples=5)
@given(
    extra_cols=st.lists(
        st.text(min_size=1, max_size=10, alphabet=st.characters(categories=("L",))),
        min_size=1,
        max_size=3,
    ),
)
def test_d1_9_validate_columns_detects_missing(extra_cols: list[str]) -> None:
    """D1-9 列校验：缺失/错误列名时能检测出缺失列。

    **Validates: Requirements 14.4**
    """
    from openpyxl import Workbook

    wb = Workbook()
    ws = wb.active

    # Remove last 2 expected columns, replace with random columns
    partial_headers = _D1_9_HEADERS[:-2] + extra_cols
    for col_idx, col_name in enumerate(partial_headers, 1):
        ws.cell(row=1, column=col_idx, value=col_name)

    missing = _validate_columns(ws, "D1-9")

    removed = set(_D1_9_HEADERS) - set(partial_headers)
    assert missing, "应检测到缺失列"
    for col in removed:
        assert col in missing, f"Expected '{col}' to be detected as missing"


# ═══════════════════════════════════════════════════════════════════════════════
# D1-13 抽样凭证核对 — 多section Round-Trip
# ═══════════════════════════════════════════════════════════════════════════════

from app.routers.wp_render_strategies._d1_import_export import (  # noqa: E402
    _export_d1_13_row,
    _export_d1_13_population_row,
    _export_d1_13_specific_row,
    _parse_d1_13_row,
    _parse_d1_13_population_row,
    _parse_d1_13_specific_row,
    _is_d1_13_section_marker,
    _D1_13_SECTION_POPULATION,
    _D1_13_SECTION_SPECIFIC,
    _D1_13_SECTION_VOUCHING,
    _D1_13_POPULATION_HEADERS,
    _D1_13_SPECIFIC_HEADERS,
    _SHEET_HEADERS,
)


_d1_13_vouching_row_st = st.fixed_dictionaries({
    "id": st.uuids().map(str),
    "seq": st.integers(min_value=1, max_value=100),
    "noteType": st.sampled_from(["银行承兑汇票", "商业承兑汇票"]),
    "noteNo": _text_st,
    "drawer": _text_st,
    "acceptor": _text_st,
    "amount": _amount_st,
    "maturityDate": _date_st,
    "existenceCheck": st.sampled_from(["已核实", "未核实", "不适用", ""]),
    "accuracyCheck": st.sampled_from(["金额一致", "金额不一致", "不适用", ""]),
    "appropriatenessCheck": st.sampled_from(["恰当", "不恰当", "不适用", ""]),
    "remark": st.text(min_size=0, max_size=20, alphabet=st.characters(categories=("L", "N"))),
    "indexRef": st.text(min_size=0, max_size=10, alphabet=st.characters(categories=("L", "N"))),
})

_d1_13_population_st = st.fixed_dictionaries({
    "populationDesc": _text_st,
    "totalCount": st.integers(min_value=0, max_value=10000),
    "totalAmount": _amount_st,
    "sampleSize": st.integers(min_value=0, max_value=500),
    "actualDrawn": st.integers(min_value=0, max_value=500),
    "sampleCalcRef": st.text(min_size=0, max_size=10, alphabet=st.characters(categories=("L", "N"))),
})


@settings(max_examples=5)
@given(population=_d1_13_population_st)
def test_d1_13_population_export_import_round_trip(population: dict) -> None:
    """D1-13 抽样总体标量 export → import round-trip preserves population fields."""
    headers = _D1_13_POPULATION_HEADERS
    exported = _export_d1_13_population_row(population)
    assert len(exported) == len(headers)

    imported = _parse_d1_13_population_row(tuple(exported), headers)

    assert imported["populationDesc"] == population["populationDesc"].strip()
    assert abs(imported["totalCount"] - float(population["totalCount"])) < 1e-6
    assert abs(imported["totalAmount"] - float(population["totalAmount"])) < 1e-6
    assert abs(imported["sampleSize"] - float(population["sampleSize"])) < 1e-6
    assert abs(imported["actualDrawn"] - float(population["actualDrawn"])) < 1e-6
    assert imported["sampleCalcRef"] == population["sampleCalcRef"].strip()


@settings(max_examples=5)
@given(rows=st.lists(_d1_13_vouching_row_st, min_size=1, max_size=5))
def test_d1_13_vouching_export_import_round_trip(rows: list[dict]) -> None:
    """D1-13 凭证核对明细 export → import round-trip preserves key fields."""
    headers = _SHEET_HEADERS["D1-13"]

    for i, original in enumerate(rows, start=1):
        exported = _export_d1_13_row(original, i)
        assert len(exported) == len(headers)

        imported = _parse_d1_13_row(tuple(exported), headers, i)

        assert imported["noteType"] == original["noteType"]
        assert imported["noteNo"] == original["noteNo"]
        assert imported["drawer"] == original["drawer"]
        assert imported["acceptor"] == original["acceptor"]
        assert abs(imported["amount"] - float(original["amount"])) < 1e-6
        assert imported["existenceCheck"] == original["existenceCheck"]
        assert imported["accuracyCheck"] == original["accuracyCheck"]
        assert imported["appropriatenessCheck"] == original["appropriatenessCheck"]
        assert imported["id"] != ""


@settings(max_examples=5)
@given(
    description=_text_st,
    amount=_amount_st,
    reason=_text_st,
)
def test_d1_13_specific_sample_round_trip(description: str, amount: float, reason: str) -> None:
    """D1-13 特定样本 description 经「票据号码」列 round-trip。"""
    headers = _D1_13_SPECIFIC_HEADERS
    original = {"description": description, "amount": amount, "reason": reason}

    exported = _export_d1_13_specific_row(original, 1)
    imported = _parse_d1_13_specific_row(tuple(exported), headers, 1)

    assert imported["description"] == description.strip()
    assert abs(imported["amount"] - amount) < 1e-6
    assert imported["reason"] == reason.strip()
    assert imported["id"] != ""


def test_d1_13_template_has_multi_section_markers() -> None:
    """D1-13 模板应含四个 section 分隔行。"""
    wb = _create_template_wb("D1-13")
    ws = wb.active

    markers = []
    for row in ws.iter_rows(values_only=True):
        first = row[0] if row else None
        if first and isinstance(first, str):
            marker = _is_d1_13_section_marker(row)
            if marker:
                markers.append(marker)

    assert markers == ["population", "specific", "vouching", "conclusion"]


def test_is_d1_13_section_marker_detects_sections() -> None:
    """_is_d1_13_section_marker 识别各 section 分隔符。"""
    assert _is_d1_13_section_marker((_D1_13_SECTION_POPULATION,)) == "population"
    assert _is_d1_13_section_marker((_D1_13_SECTION_SPECIFIC,)) == "specific"
    assert _is_d1_13_section_marker((_D1_13_SECTION_VOUCHING,)) == "vouching"
    assert _is_d1_13_section_marker(("客户A", 100)) is None


def test_d1_13_multi_section_workbook_parse() -> None:
    """构建多 section xlsx 并验证 population + vouching 行解析逻辑。"""
    from openpyxl import Workbook

    population = {
        "populationDesc": "期末在库票据",
        "totalCount": 120,
        "totalAmount": 5000000.0,
        "sampleSize": 25,
        "actualDrawn": 25,
        "sampleCalcRef": "SAM-001",
    }
    vouching = {
        "seq": 1,
        "noteType": "银行承兑汇票",
        "noteNo": "NO-001",
        "drawer": "甲公司",
        "acceptor": "某银行",
        "amount": 100000.0,
        "maturityDate": "2025-12-31",
        "existenceCheck": "已核实",
        "accuracyCheck": "金额一致",
        "appropriatenessCheck": "恰当",
        "remark": "",
        "indexRef": "D1-13-1",
    }

    wb = Workbook()
    ws = wb.active
    ws.append([_D1_13_SECTION_POPULATION])
    ws.append(_D1_13_POPULATION_HEADERS)
    ws.append(_export_d1_13_population_row(population))
    ws.append([_D1_13_SECTION_VOUCHING])
    ws.append(_SHEET_HEADERS["D1-13"])
    ws.append(_export_d1_13_row(vouching, 1))

    parsed_pop: dict = {}
    parsed_vouching: list[dict] = []
    current_section: str | None = None
    current_headers: list[str] = []
    vouching_seq = 0

    for row in ws.iter_rows(values_only=True):
        if all(v is None for v in row):
            continue
        marker = _is_d1_13_section_marker(row)
        if marker == "population":
            current_section = "population"
            current_headers = _D1_13_POPULATION_HEADERS
            continue
        if marker == "vouching":
            current_section = "vouching"
            current_headers = _SHEET_HEADERS["D1-13"]
            continue
        if not current_section:
            continue
        first = str(row[0]).strip() if row[0] is not None else ""
        if first in _D1_13_POPULATION_HEADERS or first in _SHEET_HEADERS["D1-13"]:
            current_headers = [str(c).strip() if c else "" for c in row if c is not None]
            continue
        if current_section == "population" and not parsed_pop:
            parsed_pop = _parse_d1_13_population_row(row, current_headers)
        elif current_section == "vouching":
            vouching_seq += 1
            parsed_vouching.append(_parse_d1_13_row(row, current_headers, vouching_seq))

    assert parsed_pop["populationDesc"] == population["populationDesc"]
    assert parsed_pop["totalCount"] == population["totalCount"]
    assert abs(parsed_pop["totalAmount"] - population["totalAmount"]) < 1e-6
    assert len(parsed_vouching) == 1
    assert parsed_vouching[0]["noteNo"] == vouching["noteNo"]
    assert abs(parsed_vouching[0]["amount"] - vouching["amount"]) < 1e-6

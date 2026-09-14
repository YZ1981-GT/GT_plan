"""G0A 函证程序表模板守卫

spec: g0-confirmation-source-alignment，Task 17（Property 19 / 20）

裁决者 = 源模板 `backend/wp_templates/G/G0 投资循环函证.xlsx` 的 `函证程序表G0A`。

三件事：
1. `/tables/G0A` 的 12 条 `content`/`ref_index`/`program_category` 与源模板逐字一致
2. `get_template('G0A')` 返回的是 12 条版本（而非根级 8 条自造版本）
3. `applicable_default` 分布零回归 —— 并显式记录「写 `"no"` 是死配置」的实证
"""

from __future__ import annotations

import json
from pathlib import Path

import openpyxl
import pytest

_REPO_ROOT = Path(__file__).resolve().parents[2]
_JSON_PATH = _REPO_ROOT / "backend" / "data" / "procedure_table_templates.json"
_XLSX_PATH = _REPO_ROOT / "backend" / "wp_templates" / "G" / "G0 投资循环函证.xlsx"

TABLE_CODE = "G0A"
SHEET_NAME = "函证程序表G0A"

# 改造前实测分布（2026-08-03，本 spec 未改动任何 applicable_default）
APPLICABLE_DEFAULT_BASELINE: dict[str | None, int] = {"yes": 1421, "no": 1, None: 18}
TABLES_COUNT_BASELINE = 121
ROOT_ENTRY_COUNT_BASELINE = 66

# 根级独有 → `get_template` 返回 None（运行时不可用），属平台级议题
ROOT_ONLY_DEAD_AT_RUNTIME = {"L0A", "M1A", "S1", "S2", "S3", "S8", "S10", "S11", "S13"}


def _norm_lines(text: str | None) -> str:
    """逐行去行尾空白（不动行内与行首）。"""
    if text is None:
        return ""
    return "\n".join(line.rstrip() for line in str(text).split("\n"))


@pytest.fixture(scope="module")
def data() -> dict:
    return json.loads(_JSON_PATH.read_text(encoding="utf-8"))


@pytest.fixture(scope="module")
def g0a_items(data: dict) -> list[dict]:
    return data["tables"][TABLE_CODE]["items"]


@pytest.fixture(scope="module")
def source_rows() -> list[dict]:
    """源模板 R7:R18 的 seq / content / 程序分类 / 底稿索引号。"""
    wb = openpyxl.load_workbook(_XLSX_PATH, data_only=True)
    ws = wb[SHEET_NAME]
    rows = []
    for row in range(7, 19):
        rows.append(
            {
                "seq": int(str(ws.cell(row, 1).value).strip()),
                "content": ws.cell(row, 2).value,
                "program_category": ws.cell(row, 4).value,
                "ref_index": ws.cell(row, 5).value,
            }
        )
    return rows


# ─── Property 19 —— 12 条与源模板逐字一致 + 程序分类落地 ─────────────────────


class TestG0AMatchesSource:
    def test_twelve_items(self, g0a_items):
        assert len(g0a_items) == 12
        assert [it["seq"] for it in g0a_items] == list(range(1, 13))

    def test_content_matches_source(self, g0a_items, source_rows):
        """逐行去尾空白后比对。

        JSON 侧对源模板做过行尾空白归一（实测 seq=10 源模板在
        `（1）检查交易发生的记账凭证和相关的支持性证据；` 后多一个空格），
        这是既有事实，不在本 spec 范围内回改；但**文字本身**必须逐字一致。
        """
        for item, src in zip(g0a_items, source_rows, strict=True):
            assert _norm_lines(item["content"]) == _norm_lines(src["content"]), (
                f"seq={item['seq']} content 与源模板不一致（已去行尾空白）"
            )

    def test_content_differs_from_source_only_by_trailing_space(self, g0a_items, source_rows):
        """反向自检：确认差异**只**是行尾空白，且确实存在（否则上一条退化为逐字比对）。

        若某天 JSON 与源模板出现真实文字差异，本条会把它暴露成"非空白差异"。
        """
        whitespace_only_diffs: list[int] = []
        for item, src in zip(g0a_items, source_rows, strict=True):
            if item["content"] == src["content"]:
                continue
            assert _norm_lines(item["content"]) == _norm_lines(src["content"]), (
                f"seq={item['seq']} 存在**非空白**文字差异，须人工裁决"
            )
            whitespace_only_diffs.append(item["seq"])
        assert whitespace_only_diffs == [10], (
            f"行尾空白差异清单变化：{whitespace_only_diffs}（原为 [10]）"
        )

    def test_ref_index_matches_source(self, g0a_items, source_rows):
        """源模板空格 → JSON 里是空串（不是 None）。"""
        for item, src in zip(g0a_items, source_rows, strict=True):
            expected = src["ref_index"] or ""
            assert item["ref_index"] == expected, f"seq={item['seq']} ref_index 与源模板不一致"

    def test_two_items_have_no_ref_index(self, g0a_items):
        assert {it["seq"] for it in g0a_items if not it["ref_index"]} == {8, 12}

    def test_program_category_matches_source(self, g0a_items, source_rows):
        for item, src in zip(g0a_items, source_rows, strict=True):
            assert item.get("program_category") == src["program_category"], (
                f"seq={item['seq']} program_category 缺失或与源模板不一致 —— "
                "跑 `python backend/scripts/fix/fix_g0a_program_category.py --apply`"
            )

    def test_all_categories_non_empty(self, g0a_items):
        assert all(it.get("program_category") for it in g0a_items)

    def test_category_distribution(self, g0a_items):
        from collections import Counter

        got = Counter(it["program_category"] for it in g0a_items)
        assert got == Counter(
            {
                "常规★": 8,
                "备选": 2,
                "IPO/上市/新三板/重组/舞弊应对": 1,
                "舞弊应对/IPO/上市/新三板/重组": 1,
            }
        )

    def test_optional_and_ipo_seq(self, g0a_items):
        by_seq = {it["seq"]: it["program_category"] for it in g0a_items}
        assert {s for s, c in by_seq.items() if c == "备选"} == {7, 8}
        assert {s for s, c in by_seq.items() if "IPO" in c} == {2, 11}

    def test_category_vocabulary_has_platform_precedent(self, data, g0a_items):
        """`IPO/上市/新三板/重组/舞弊应对` 与 D4-22A 既有取值逐字相同 → 词表非自造。"""
        existing = {
            it.get("program_category")
            for tbl in data["tables"].values()
            for it in tbl.get("items", [])
            if it.get("program_category")
        }
        assert "IPO/上市/新三板/重组/舞弊应对" in existing
        # D4-22A 是先例来源
        d4 = data["tables"].get("D4-22A", {}).get("items", [])
        assert any(it.get("program_category") == "IPO/上市/新三板/重组/舞弊应对" for it in d4)


# ─── Property 19 —— 前端/render 侧确有消费方（防死字段） ──────────────────────


class TestProgramCategoryHasConsumer:
    def test_render_strategy_maps_program_category(self):
        src = (
            _REPO_ROOT
            / "backend" / "app" / "routers" / "wp_render_strategies" / "_a_program.py"
        ).read_text(encoding="utf-8")
        assert '"program_category": it.get("category") or it.get("program_category")' in src

    def test_frontend_renders_category_column_and_filter(self):
        src = (
            _REPO_ROOT
            / "audit-platform" / "frontend" / "src" / "components" / "workpaper"
            / "GtAProgramConsole.vue"
        ).read_text(encoding="utf-8")
        assert 'prop="program_category"' in src, "前端「类别」列消失 → program_category 变成死字段"
        assert "activeCategory" in src, "类别筛选消失"
        assert "hasCategory" in src, "空分类隐藏列的逻辑消失"
        assert "categoryTagType" in src


# ─── Property 20 —— 运行时权威唯一 + 根级分叉可见 + 零回归 ───────────────────


class TestRuntimeAuthority:
    def test_get_template_returns_tables_version(self):
        from app.services.procedure_table_auto_service import get_template

        tpl = get_template(TABLE_CODE)
        assert tpl is not None
        assert tpl["name"] == "函证程序表"
        assert len(tpl["items"]) == 12, "get_template 应返回 tables 版本（12 条）而非根级 8 条"

    def test_root_level_g0a_is_divergent_dead_data(self, data):
        """根级 `/G0A` 与 `tables` 版本内容分叉 —— 平台级议题，本 spec 不删。

        🔴 若这条断言打红（两者内容一致或根级已删除），说明平台级收敛已发生，
        须复核 g0-confirmation-source-alignment 的 Requirement 8.4 与遗留登记。
        平台级议题全貌：顶层共 66 条同族条目，57 条与 `tables` 重复且内容分叉，
        另 9 条根级独有在运行时不可用（见 test_root_only_tables_are_dead_at_runtime）。
        """
        root = data.get(TABLE_CODE)
        if root is None:
            pytest.skip("根级 /G0A 已被平台级 spec 移除 —— 请复核 R8.4 与遗留登记")
        assert len(root["items"]) != len(data["tables"][TABLE_CODE]["items"])
        assert root.get("name") != data["tables"][TABLE_CODE].get("name")

    def test_root_only_tables_are_dead_at_runtime(self, data):
        """9 条根级独有的程序表 `get_template` 返回 None —— 平台级缺陷，只登记不修。"""
        from app.services.procedure_table_auto_service import get_template

        tables = set(data["tables"])
        root = {k for k in data if k not in ("version", "description", "tables")}
        assert root - tables == ROOT_ONLY_DEAD_AT_RUNTIME
        for code in sorted(ROOT_ONLY_DEAD_AT_RUNTIME):
            assert get_template(code) is None, (
                f"{code} 已可用 → 平台级收敛已发生，请复核遗留登记"
            )

    def test_structure_counts_unchanged(self, data):
        assert len(data["tables"]) == TABLES_COUNT_BASELINE
        root = {k for k in data if k not in ("version", "description", "tables")}
        assert len(root) == ROOT_ENTRY_COUNT_BASELINE

    def test_applicable_default_distribution_unchanged(self, data):
        """本 spec 未改动任何 `applicable_default` —— 分布须与改造前实测一致。"""
        from collections import Counter

        got = Counter(
            it.get("applicable_default")
            for tbl in data["tables"].values()
            for it in tbl.get("items", [])
        )
        assert dict(got) == APPLICABLE_DEFAULT_BASELINE

    def test_g0a_applicable_default_all_yes(self, g0a_items):
        """🔴 实证：`applicable_default` 写 `"no"` 是死配置。

        `_check_applicable` 在无 `applicable_categories` 时原样返回 `applicable_default`，
        而 `_a_program.py` 只把 `applicable == "na"` 映射成 `status='not_applicable'`
        → `"no"` 与 `"yes"` 渲染结果相同。故「备选/IPO 专项默认不执行」不能靠它表达，
        改由 `program_category` + 前端类别筛选 + 批量裁剪落地（Requirement 8.2/8.3）。
        """
        assert {it["applicable_default"] for it in g0a_items} == {"yes"}

    def test_render_only_treats_na_as_not_applicable(self):
        src = (
            _REPO_ROOT
            / "backend" / "app" / "routers" / "wp_render_strategies" / "_a_program.py"
        ).read_text(encoding="utf-8")
        assert 'status = "not_applicable" if applicable == "na" else "pending"' in src, (
            "该映射若改变，`applicable_default: \"no\"` 可能不再是死配置 → "
            "复核 Requirement 8.3 与本 spec 的遗留登记"
        )


# ─── 幂等脚本自身 ────────────────────────────────────────────────────────────


class TestFixScript:
    def test_check_reports_no_debt(self):
        import subprocess
        import sys

        script = _REPO_ROOT / "backend" / "scripts" / "fix" / "fix_g0a_program_category.py"
        proc = subprocess.run(
            [sys.executable, str(script), "--check"],
            capture_output=True,
            text=True,
            encoding="utf-8",
            errors="replace",
            cwd=str(_REPO_ROOT),
        )
        assert proc.returncode == 0, f"--check 有欠账:\n{proc.stdout}\n{proc.stderr}"

    def test_constants_cross_locked_with_source(self, source_rows):
        """脚本常量表与源模板双向锁死。"""
        from importlib.util import module_from_spec, spec_from_file_location

        path = _REPO_ROOT / "backend" / "scripts" / "fix" / "fix_g0a_program_category.py"
        spec = spec_from_file_location("_fix_g0a", path)
        assert spec and spec.loader
        mod = module_from_spec(spec)
        spec.loader.exec_module(mod)

        expected = {r["seq"]: r["program_category"] for r in source_rows}
        assert mod.PROGRAM_CATEGORY == expected


# ─── 反向自检 ────────────────────────────────────────────────────────────────


class TestReverseSelfChecks:
    def test_selfcheck_source_read_is_live(self, source_rows):
        assert len(source_rows) == 12
        assert all(r["content"] for r in source_rows)
        assert source_rows[0]["program_category"] == "常规★"

    def test_selfcheck_wrong_category_would_fail(self, g0a_items):
        """若把「备选」误写成「常规★」，`test_category_distribution` 必打红。"""
        from collections import Counter

        tampered = Counter(["常规★"] * 10 + ["IPO/上市/新三板/重组/舞弊应对", "舞弊应对/IPO/上市/新三板/重组"])
        got = Counter(it["program_category"] for it in g0a_items)
        assert got != tampered

"""章节增量刷新就地替换与人工编辑检测测试 — Task 7.1

覆盖：
- Property 7：人工编辑检测三态（未改 ⇒ False / 改一字 ⇒ True / 基线 NULL ⇒ False）
- Property 8：刷新后该章节**位置不变**，其余章节文字逐字不变
- 批量刷新在同一份 doc 上替换后只落一个版本（修既有缺陷：逐章节从原始字节重解析，
  上一章节的替换被丢弃 → 只有最后一章生效）
- 反向自检：拿 source_snapshot_hash 当基线（旧实现）必然判「有人工编辑」
"""

from __future__ import annotations

import asyncio
import uuid
from io import BytesIO

import pytest
import sqlalchemy as sa
from docx import Document
from sqlalchemy.dialects.sqlite.base import SQLiteTypeCompiler
from sqlalchemy.ext.asyncio import async_sessionmaker, create_async_engine

from app.models.audit_platform_models import DeliverableSectionState, TrialBalance
from app.models.phase13_models import WordExportTask
from app.models.phase14_models import TraceEvent
from app.models.report_models import DisclosureNote
from app.services.deliverable_refresh_service import DeliverableRefreshService
from app.services.section_anchor_utils import (
    SectionBlock,
    block_text_hash,
    resolve_section_blocks,
    scan_anchor_blocks,
    write_section_anchors,
)
from app.services.word_doc_utils import remove_section_markers, scan_section_blocks

SQLiteTypeCompiler.visit_JSONB = SQLiteTypeCompiler.visit_JSON

PROJECT_ID = uuid.uuid4()
TASK_ID = uuid.uuid4()
ACTOR_ID = uuid.uuid4()
YEAR = 2025

_TABLES = [
    DeliverableSectionState.__table__,
    TrialBalance.__table__,
    DisclosureNote.__table__,
    WordExportTask.__table__,
    TraceEvent.__table__,
]


def _delivered_doc(sections: dict[str, list[str]]) -> Document:
    doc = Document()
    doc.add_paragraph("交付文档首段")
    for code, lines in sections.items():
        doc.add_paragraph(f"##SECTION:{code}##")
        for line in lines:
            doc.add_paragraph(line)
        doc.add_paragraph(f"##/SECTION:{code}##")
    blocks = {b.section_code: b for b in scan_section_blocks(doc)}
    write_section_anchors(
        doc,
        [
            SectionBlock(
                section_code=c, open_el=blocks[c].open_el, close_el=blocks[c].close_el
            )
            for c in sections
        ],
    )
    remove_section_markers(doc)
    return doc


def _to_bytes(doc: Document) -> bytes:
    buf = BytesIO()
    doc.save(buf)
    return buf.getvalue()


def _all_paragraph_texts(docx_bytes: bytes) -> list[str]:
    from docx.oxml.ns import qn

    doc = Document(BytesIO(docx_bytes))
    ns_w = "http://schemas.openxmlformats.org/wordprocessingml/2006/main"
    return [
        "".join(t.text or "" for t in p.iter(f"{{{ns_w}}}t"))
        for p in doc.element.body.iter(qn("w:p"))
    ]


def _run(scenario):
    async def _main():
        engine = create_async_engine("sqlite+aiosqlite:///:memory:", echo=False)
        async with engine.begin() as conn:
            await conn.run_sync(
                DeliverableSectionState.metadata.create_all, tables=_TABLES
            )
        factory = async_sessionmaker(engine, expire_on_commit=False)
        try:
            async with factory() as session:
                return await scenario(session)
        finally:
            await engine.dispose()

    return asyncio.run(_main())


async def _setup(session, *, notes: dict[str, str], rendered: dict[str, str | None]):
    session.add(
        WordExportTask(
            id=TASK_ID,
            project_id=PROJECT_ID,
            doc_type="disclosure_notes",
            status="editing",
            created_by=ACTOR_ID,
        )
    )
    for code, text in notes.items():
        session.add(
            DisclosureNote(
                project_id=PROJECT_ID,
                year=YEAR,
                note_section=code,
                section_title=code,
                text_content=text,
                is_deleted=False,
            )
        )
    for code, block_hash in rendered.items():
        session.add(
            DeliverableSectionState(
                word_export_task_id=TASK_ID,
                project_id=PROJECT_ID,
                year=YEAR,
                section_code=code,
                source_snapshot_hash="0" * 64,
                rendered_block_hash=block_hash,
                is_stale=True,
            )
        )
    await session.flush()


# ─── Property 7：人工编辑检测三态 ────────────────────────────────────────────


@pytest.mark.parametrize(
    "case,doc_lines,baseline_from_doc,expect_edits",
    [
        ("未改", ["原始内容"], True, False),
        ("改一字", ["原始内容"], "mutated", True),
        ("基线NULL", ["原始内容"], None, False),
    ],
)
def test_property_7_user_edit_detection_three_states(
    case, doc_lines, baseline_from_doc, expect_edits
):
    # Feature: deliverable-lineage-wiring-and-writeback-closure, Property 7
    doc = _delivered_doc({"八、1": doc_lines})
    block = scan_anchor_blocks(doc)[0]
    actual_hash = block_text_hash(block.elements)

    if baseline_from_doc is True:
        baseline = actual_hash
    elif baseline_from_doc == "mutated":
        # 基线记录的是「原始内容」，而 doc 里已被人工改成别的字
        mutated = _delivered_doc({"八、1": ["原始内容被人改过"]})
        baseline = actual_hash
        doc = mutated
        block = scan_anchor_blocks(doc)[0]
    else:
        baseline = None

    async def _scenario(session):
        await _setup(
            session, notes={"八、1": "DB 文字"}, rendered={"八、1": baseline}
        )
        svc = DeliverableRefreshService(session)
        return await svc._detect_user_edits(
            TASK_ID, PROJECT_ID, YEAR, "八、1", block
        )

    assert _run(_scenario) is expect_edits, f"用例「{case}」判定错误"


def test_reverse_selfcheck_snapshot_hash_as_baseline_always_reports_edits():
    """反向自检：用 source_snapshot_hash 当基线（旧实现）必判「有人工编辑」。

    钉住「两个哈希域不可互换」：source_snapshot_hash 是 DB 源数据域，
    与块内渲染文字域永远不等 ⇒ 刷新恒 requires_confirm。
    """
    from app.services.deliverable_section_state_service import (
        compute_snapshot_hash_from_parts,
    )

    doc = _delivered_doc({"八、1": ["一致的内容"]})
    block = scan_anchor_blocks(doc)[0]
    snapshot_hash = compute_snapshot_hash_from_parts(
        section_code="八、1",
        text_content="一致的内容",
        table_data=None,
        audited_amounts=[],
    )
    # 即便块内文字与 text_content 完全一致，两个哈希也不等
    assert block_text_hash(block.elements) != snapshot_hash


# ─── Property 8：刷新保位 ────────────────────────────────────────────────────


def test_property_8_refresh_keeps_position_and_other_sections_intact():
    # Feature: deliverable-lineage-wiring-and-writeback-closure, Property 8
    """刷新中间章节：位置序号不变，其余章节文字逐字不变。"""
    doc = _delivered_doc(
        {"八、1": ["甲原文"], "八、2": ["乙原文"], "八、3": ["丙原文"]}
    )
    original_bytes = _to_bytes(doc)
    before = _all_paragraph_texts(original_bytes)
    assert before.index("乙原文") == 2  # 首段 + 甲 + 乙 + 丙

    async def _scenario(session):
        blk = {b.section_code: b for b in scan_anchor_blocks(Document(BytesIO(original_bytes)))}
        await _setup(
            session,
            notes={"八、1": "甲原文", "八、2": "乙的新内容", "八、3": "丙原文"},
            rendered={c: block_text_hash(b.elements) for c, b in blk.items()},
        )
        svc = DeliverableRefreshService(session)
        stored: dict = {}

        async def _fake_store(_self, task_id, **kw):
            stored["bytes"] = kw["docx_bytes"]

            class _V:
                version_no = 7

            class _R:
                version = _V()

            return _R()

        from app.services import deliverable_service as ds_mod

        original = ds_mod.DeliverableService.render_and_store
        ds_mod.DeliverableService.render_and_store = _fake_store  # type: ignore[assignment]
        try:
            result = await svc.refresh_section(
                TASK_ID, PROJECT_ID, YEAR, "八、2", ACTOR_ID
            , docx_bytes=original_bytes)
        finally:
            ds_mod.DeliverableService.render_and_store = original  # type: ignore[assignment]
        return result, stored.get("bytes")

    result, new_bytes = _run(_scenario)

    assert result["refreshed"] == ["八、2"], result
    assert new_bytes is not None
    after = _all_paragraph_texts(new_bytes)

    # 位置不变：新内容仍在第 3 段（index 2），不是被追加到文末
    assert after[2] == "乙的新内容", after
    assert after[-1] != "乙的新内容", "刷新内容被追加到文档末尾（Property 8 违反）"
    # 其余章节逐字不变
    assert after[1] == "甲原文"
    assert after[3] == "丙原文"
    assert len(after) == len(before)

    # 锚点仍在，区间仍能定位到刷新后的内容
    doc2 = Document(BytesIO(new_bytes))
    blocks2 = {b.section_code: b for b in scan_anchor_blocks(doc2)}
    assert set(blocks2) == {"八、1", "八、2", "八、3"}
    from app.services.section_anchor_utils import block_text_of

    assert block_text_of(blocks2["八、2"].elements) == "乙的新内容"


def test_refresh_updates_rendered_block_hash_and_clears_stale():
    """刷新后 rendered_block_hash 更新为新内容哈希、is_stale 清除。"""
    doc = _delivered_doc({"八、1": ["旧内容"]})
    original_bytes = _to_bytes(doc)

    async def _scenario(session):
        blk = scan_anchor_blocks(Document(BytesIO(original_bytes)))[0]
        await _setup(
            session,
            notes={"八、1": "新内容"},
            rendered={"八、1": block_text_hash(blk.elements)},
        )
        svc = DeliverableRefreshService(session)

        async def _fake_store(_self, task_id, **kw):
            class _V:
                version_no = 2

            class _R:
                version = _V()

            return _R()

        from app.services import deliverable_service as ds_mod

        original = ds_mod.DeliverableService.render_and_store
        ds_mod.DeliverableService.render_and_store = _fake_store  # type: ignore[assignment]
        try:
            await svc.refresh_section(
                TASK_ID, PROJECT_ID, YEAR, "八、1", ACTOR_ID,
                docx_bytes=original_bytes,
            )
        finally:
            ds_mod.DeliverableService.render_and_store = original  # type: ignore[assignment]

        row = (
            await session.execute(
                sa.select(DeliverableSectionState).where(
                    DeliverableSectionState.section_code == "八、1"
                )
            )
        ).scalar_one()
        return row.rendered_block_hash, row.is_stale

    new_hash, is_stale = _run(_scenario)

    # 期望值 = 「新内容」单段落的规范化哈希（与导出侧同一函数口径）
    import hashlib

    expected = hashlib.sha256("新内容".encode("utf-8")).hexdigest()
    assert is_stale is False
    assert new_hash == expected, (
        f"刷新后 rendered_block_hash 未更新为新内容哈希: {new_hash} != {expected}"
    )


# ─── 批量刷新只落一个版本、不丢改动 ─────────────────────────────────────────


def test_batch_refresh_applies_all_sections_in_one_version():
    """批量刷新：所有 stale 章节都生效，且只产生**一个**新版本。

    旧实现逐章节调 refresh_section 并每次传同一份原始字节 → 每次从原始字节重解析，
    上一章节的替换被丢弃 → 最终版本只含最后一章，且产生 N 个版本。
    """
    doc = _delivered_doc({"八、1": ["甲旧"], "八、2": ["乙旧"]})
    original_bytes = _to_bytes(doc)

    async def _scenario(session):
        blk = {
            b.section_code: b
            for b in scan_anchor_blocks(Document(BytesIO(original_bytes)))
        }
        await _setup(
            session,
            notes={"八、1": "甲新", "八、2": "乙新"},
            rendered={c: block_text_hash(b.elements) for c, b in blk.items()},
        )
        svc = DeliverableRefreshService(session)
        calls: list[bytes] = []

        async def _fake_store(_self, task_id, **kw):
            calls.append(kw["docx_bytes"])

            class _V:
                version_no = len(calls)

            class _R:
                version = _V()

            return _R()

        from app.services import deliverable_service as ds_mod

        original = ds_mod.DeliverableService.render_and_store
        ds_mod.DeliverableService.render_and_store = _fake_store  # type: ignore[assignment]
        try:
            result = await svc.refresh_all_stale_sections(
                TASK_ID, PROJECT_ID, YEAR, ACTOR_ID, docx_bytes=original_bytes
            )
        finally:
            ds_mod.DeliverableService.render_and_store = original  # type: ignore[assignment]
        return result, calls

    result, calls = _run(_scenario)

    assert sorted(result["refreshed"]) == ["八、1", "八、2"]
    assert len(calls) == 1, f"批量刷新应只落一个版本，实际 {len(calls)} 个"
    texts = _all_paragraph_texts(calls[0])
    assert "甲新" in texts and "乙新" in texts, texts
    assert "甲旧" not in texts and "乙旧" not in texts


def test_batch_refresh_pending_confirm_writes_nothing():
    """有待确认章节且未确认 ⇒ 整批不写入（不产生半写状态）。"""
    doc = _delivered_doc({"八、1": ["甲旧"], "八、2": ["乙旧"]})
    original_bytes = _to_bytes(doc)

    async def _scenario(session):
        # 八、2 的基线故意给错 → 判为「有人工编辑」
        blk = {
            b.section_code: b
            for b in scan_anchor_blocks(Document(BytesIO(original_bytes)))
        }
        await _setup(
            session,
            notes={"八、1": "甲新", "八、2": "乙新"},
            rendered={
                "八、1": block_text_hash(blk["八、1"].elements),
                "八、2": "f" * 64,
            },
        )
        svc = DeliverableRefreshService(session)
        calls: list[bytes] = []

        async def _fake_store(_self, task_id, **kw):  # pragma: no cover - 不应被调用
            calls.append(kw["docx_bytes"])
            raise AssertionError("待确认时不应落版本")

        from app.services import deliverable_service as ds_mod

        original = ds_mod.DeliverableService.render_and_store
        ds_mod.DeliverableService.render_and_store = _fake_store  # type: ignore[assignment]
        try:
            result = await svc.refresh_all_stale_sections(
                TASK_ID, PROJECT_ID, YEAR, ACTOR_ID, docx_bytes=original_bytes
            )
        finally:
            ds_mod.DeliverableService.render_and_store = original  # type: ignore[assignment]
        return result, calls

    result, calls = _run(_scenario)
    assert result["requires_confirm"] is True
    assert result["pending_confirm_sections"] == ["八、2"]
    assert result["refreshed"] == []
    assert calls == []


# ─── 降级：无锚点无标记 ──────────────────────────────────────────────────────


def test_refresh_on_legacy_docx_skips_without_error():
    doc = Document()
    doc.add_paragraph("没有章节标识")
    legacy = _to_bytes(doc)

    async def _scenario(session):
        await _setup(session, notes={"八、1": "新"}, rendered={"八、1": None})
        svc = DeliverableRefreshService(session)
        mode, _ = resolve_section_blocks(Document(BytesIO(legacy)))
        result = await svc.refresh_section(
            TASK_ID, PROJECT_ID, YEAR, "八、1", ACTOR_ID, docx_bytes=legacy
        )
        return mode, result

    mode, result = _run(_scenario)
    assert mode == "none"
    assert result["skipped"] == ["八、1"]
    assert result["version_no"] is None

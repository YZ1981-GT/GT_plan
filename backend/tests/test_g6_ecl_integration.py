"""G6 其他债权投资(ECL组) — 集成测试.

Spec: .kiro/specs/g6-other-bond-investment-ecl/ Task 11.3
Validates: Requirements 1.4, 6.2, 6.3, 6.4

Tests:
  1. selfLoad: render-config端点返回正确结构 (component_type, sheets=5)
  2. 导入导出: 多表×端点 + G6-14/15 结构 roundtrip
  3. AI接口: 5个section各返回合理response structure
  4. 版本链: render策略返回responses_snapshot
"""

from __future__ import annotations

import io
from unittest.mock import AsyncMock, patch

import pytest
from httpx import ASGITransport, AsyncClient
from openpyxl import Workbook, load_workbook

from app.core.database import get_db
from app.deps import get_current_user
from app.main import app
from app.models.base import UserRole


class _FakeUser:
    id = "test-user-id"
    name = "Test User"
    email = "test@example.com"
    role = UserRole.admin


@pytest.fixture(autouse=True)
def override_deps():
    """Override DB and auth dependencies for all tests."""
    mock_db = AsyncMock()
    mock_db.commit = AsyncMock()

    # Default: execute returns empty result set
    mock_result = AsyncMock()
    mock_result.scalar_one_or_none = AsyncMock(return_value=None)
    mock_result.fetchone = AsyncMock(return_value=None)
    mock_result.fetchall = AsyncMock(return_value=[])
    mock_db.execute = AsyncMock(return_value=mock_result)

    async def _override_db():
        yield mock_db

    app.dependency_overrides[get_current_user] = lambda: _FakeUser()
    app.dependency_overrides[get_db] = _override_db
    yield mock_db
    app.dependency_overrides.clear()


def _make_xlsx_bytes(headers: list[str], rows: list[list]) -> bytes:
    """Create minimal xlsx in-memory with given headers and data rows."""
    wb = Workbook()
    ws = wb.active
    ws.append(headers)
    for row in rows:
        ws.append(row)
    buf = io.BytesIO()
    wb.save(buf)
    buf.seek(0)
    return buf.getvalue()


def _make_xlsx_with_title(title: str, headers: list[str], rows: list[list]) -> bytes:
    """Create xlsx with title row (row 1) + header row (row 2) + data rows (row 3+)."""
    wb = Workbook()
    ws = wb.active
    ws.append([title])  # row 1: title
    ws.append(headers)  # row 2: headers
    for row in rows:
        ws.append(row)
    buf = io.BytesIO()
    wb.save(buf)
    buf.seek(0)
    return buf.getvalue()


# ═══════════════════════════════════════════════════════════════════════════════
# 1. selfLoad: render-config 端点返回正确结构
# ═══════════════════════════════════════════════════════════════════════════════


class TestRenderConfig:
    """验证 render 策略返回结构：component_type + 5 sheets."""

    def test_g6_ecl_sheets_constant_has_5_sheets(self):
        """G6_ECL_SHEETS 常量包含 5 个 sheet 配置."""
        from app.routers.wp_render_strategies._g6_other_bond_investment_ecl import G6_ECL_SHEETS

        assert len(G6_ECL_SHEETS) == 5

    def test_g6_ecl_sheets_correct_codes(self):
        """5 个 sheet 的 code 字段正确."""
        from app.routers.wp_render_strategies._g6_other_bond_investment_ecl import G6_ECL_SHEETS

        expected_codes = {"G6-11", "G6-12", "G6-13", "G6-14", "G6-15"}
        actual_codes = {s["code"] for s in G6_ECL_SHEETS}
        assert actual_codes == expected_codes

    def test_g6_ecl_sheets_all_same_component_type(self):
        """每个 sheet 的 componentType 都是 g6-other-bond-investment-ecl."""
        from app.routers.wp_render_strategies._g6_other_bond_investment_ecl import G6_ECL_SHEETS

        for sheet in G6_ECL_SHEETS:
            assert sheet["componentType"] == "g6-other-bond-investment-ecl"
            assert "sheetName" in sheet
            assert "group" in sheet

    def test_g6_ecl_sheets_group_assignment(self):
        """G6-11~G6-14 属于 impairment 组，G6-15 属于 voucher 组."""
        from app.routers.wp_render_strategies._g6_other_bond_investment_ecl import G6_ECL_SHEETS

        for sheet in G6_ECL_SHEETS:
            if sheet["code"] == "G6-15":
                assert sheet["group"] == "voucher"
            else:
                assert sheet["group"] == "impairment"

    @pytest.mark.asyncio
    async def test_render_function_returns_correct_structure(self):
        """render 策略函数返回正确顶层结构（含 responses_snapshot）."""
        from app.routers.wp_render_strategies._g6_other_bond_investment_ecl import render
        from app.routers.wp_render_strategies._context import RenderContext
        from unittest.mock import MagicMock
        from uuid import uuid4

        mock_db = AsyncMock()
        mock_result = AsyncMock()
        mock_result.fetchall = AsyncMock(return_value=[])
        mock_result.fetchone = AsyncMock(return_value=None)
        mock_db.execute = AsyncMock(return_value=mock_result)

        # Create a minimal mock RenderContext
        ctx = MagicMock(spec=RenderContext)
        ctx.wp_id = uuid4()
        ctx.project_id = uuid4()
        ctx.db = mock_db

        result = await render(ctx)

        assert result is not None
        assert result["component_type"] == "g6-other-bond-investment-ecl"
        assert "sheets" in result
        assert len(result["sheets"]) == 5
        assert "responses_snapshot" in result
        assert isinstance(result["responses_snapshot"], dict)


# ═══════════════════════════════════════════════════════════════════════════════
# 2. 导入导出: 3张表×3端点 roundtrip
# ═══════════════════════════════════════════════════════════════════════════════


class TestImportExport:
    """验证 3 张表(G6-12/G6-14/G6-15)的导入导出端点."""

    # --- Export Template ---

    @pytest.mark.asyncio
    async def test_export_template_g6_12(self):
        transport = ASGITransport(app=app)
        async with AsyncClient(transport=transport, base_url="http://test") as client:
            resp = await client.post(
                "/api/workpapers/test-wp/g6-ecl/export-template?sheet=G6-12"
            )
        assert resp.status_code == 200
        assert "application/vnd.openxmlformats" in resp.headers["content-type"]

    @pytest.mark.asyncio
    async def test_export_template_g6_14(self):
        transport = ASGITransport(app=app)
        async with AsyncClient(transport=transport, base_url="http://test") as client:
            resp = await client.post(
                "/api/workpapers/test-wp/g6-ecl/export-template?sheet=G6-14"
            )
        assert resp.status_code == 200
        assert "application/vnd.openxmlformats" in resp.headers["content-type"]

    @pytest.mark.asyncio
    async def test_export_template_g6_15(self):
        transport = ASGITransport(app=app)
        async with AsyncClient(transport=transport, base_url="http://test") as client:
            resp = await client.post(
                "/api/workpapers/test-wp/g6-ecl/export-template?sheet=G6-15"
            )
        assert resp.status_code == 200
        assert "application/vnd.openxmlformats" in resp.headers["content-type"]

    # --- Export Data (empty) ---

    @pytest.mark.asyncio
    async def test_export_data_g6_12_empty(self):
        transport = ASGITransport(app=app)
        async with AsyncClient(transport=transport, base_url="http://test") as client:
            resp = await client.post(
                "/api/workpapers/test-wp/g6-ecl/export-data?sheet=G6-12"
            )
        assert resp.status_code == 200
        assert "application/vnd.openxmlformats" in resp.headers["content-type"]

    @pytest.mark.asyncio
    async def test_export_data_g6_14_empty(self):
        transport = ASGITransport(app=app)
        async with AsyncClient(transport=transport, base_url="http://test") as client:
            resp = await client.post(
                "/api/workpapers/test-wp/g6-ecl/export-data?sheet=G6-14"
            )
        assert resp.status_code == 200
        assert "application/vnd.openxmlformats" in resp.headers["content-type"]

    @pytest.mark.asyncio
    async def test_export_data_g6_15_empty(self):
        transport = ASGITransport(app=app)
        async with AsyncClient(transport=transport, base_url="http://test") as client:
            resp = await client.post(
                "/api/workpapers/test-wp/g6-ecl/export-data?sheet=G6-15"
            )
        assert resp.status_code == 200
        assert "application/vnd.openxmlformats" in resp.headers["content-type"]

    # --- Import Data ---

    @pytest.mark.asyncio
    async def test_import_data_g6_12(self):
        """Import G6-12 xlsx with correct headers → 200, imported_count > 0."""
        xlsx_bytes = _make_xlsx_bytes(
            ["投资项目", "摊余成本余额①", "公允价值", "预期信用损失率②",
             "坏账准备③", "账面价值④", "余额调整⑤",
             "调整后损失率②A", "坏账调整⑥", "阶段", "OCI影响",
             "审定余额⑦", "审定坏账⑧", "审定账面价值⑨",
             "审定公允价值", "上年坏账", "本年计提",
             "本年转回", "OCI调整", "差异说明", "索引"],
            [
                ["国开债2024", 10000000, 10200000, 0.01, 100000, 9900000, 200000,
                 0.02, 0, "Stage1", 0, 10200000, 100000, 10100000,
                 10300000, 80000, 20000, 0, 0, "", ""],
            ],
        )
        transport = ASGITransport(app=app)
        async with AsyncClient(transport=transport, base_url="http://test") as client:
            resp = await client.post(
                "/api/workpapers/test-wp/g6-ecl/import-data?sheet=G6-12",
                files={"file": ("test.xlsx", xlsx_bytes, "application/vnd.openxmlformats-officedocument.spreadsheetml.sheet")},
            )
        assert resp.status_code == 200
        data = resp.json()
        payload = data.get("data", data)
        assert payload["imported_count"] > 0

    @pytest.mark.asyncio
    async def test_export_template_g6_14_has_dual_tabs(self):
        transport = ASGITransport(app=app)
        async with AsyncClient(transport=transport, base_url="http://test") as client:
            resp = await client.post(
                "/api/workpapers/test-wp/g6-ecl/export-template?sheet=G6-14"
            )
        assert resp.status_code == 200
        wb = load_workbook(io.BytesIO(resp.content))
        assert any("转回检查" in n for n in wb.sheetnames)
        assert any("核销检查" in n for n in wb.sheetnames)

    @pytest.mark.asyncio
    async def test_import_data_g6_14_dual_tabs(self):
        """Import G6-14 双 Tab xlsx → imported_count >= 2."""
        from app.routers.wp_render_strategies._g6_other_bond_investment_ecl_import_export import (
            _build_g6_14_multi_sheet_workbook,
        )
        wb = _build_g6_14_multi_sheet_workbook(
            [{
                "seq": 1, "unitName": "农发债", "kind": "转回",
                "reversalReason": "改善", "recoveryMethod": "",
                "originalBasis": "", "reversalAmount": 1000,
                "accumulatedProvision": 5000, "reasonAnalysis": "ok",
                "isReasonable": "合理", "indexRef": "",
            }],
            [{
                "seq": 1, "unitName": "企债", "writeOffType": "债",
                "writeOffAmount": 2000, "writeOffReason": "破产",
                "writeOffProcedure": "董事会", "isRelatedParty": False,
                "reasonAnalysis": "", "isReasonable": "合理", "indexRef": "",
            }],
        )
        buf = io.BytesIO()
        wb.save(buf)
        transport = ASGITransport(app=app)
        async with AsyncClient(transport=transport, base_url="http://test") as client:
            resp = await client.post(
                "/api/workpapers/test-wp/g6-ecl/import-data?sheet=G6-14",
                files={"file": ("g6-14.xlsx", buf.getvalue(),
                       "application/vnd.openxmlformats-officedocument.spreadsheetml.sheet")},
            )
        assert resp.status_code == 200
        payload = resp.json().get("data", resp.json())
        assert payload["imported_count"] >= 2

    @pytest.mark.asyncio
    async def test_import_data_g6_14(self):
        """Import G6-14 xlsx → 200, imported_count > 0."""
        # G6-14 uses parse_upload_xlsx with header_row=2,
        # so row 1 = title, row 2 = headers, row 3+ = data
        xlsx_bytes = _make_xlsx_with_title(
            "G6-14 减值准备转回（收回）、核销检查表",
            ["序号", "投资项目", "转回/核销类型", "金额",
             "原因", "审批程序", "合理性结论", "索引"],
            [
                [1, "农发债2024", "转回", 50000, "信用风险改善", "经理审批", "合理", ""],
                [2, "企业债2023", "核销", 200000, "债务人破产", "董事会", "合理", ""],
            ],
        )
        transport = ASGITransport(app=app)
        async with AsyncClient(transport=transport, base_url="http://test") as client:
            resp = await client.post(
                "/api/workpapers/test-wp/g6-ecl/import-data?sheet=G6-14",
                files={"file": ("test.xlsx", xlsx_bytes, "application/vnd.openxmlformats-officedocument.spreadsheetml.sheet")},
            )
        assert resp.status_code == 200
        data = resp.json()
        payload = data.get("data", data)
        assert payload["imported_count"] >= 2

    @pytest.mark.asyncio
    async def test_import_data_g6_15(self):
        """Import G6-15 xlsx with correct headers → 200."""
        xlsx_bytes = _make_xlsx_bytes(
            ["日期", "凭证号", "业务内容", "对方科目",
             "明细", "借方", "贷方", "附件",
             "支持性文件",
             "核对1-原始凭证(是/否)", "核对2-授权(是/否)",
             "核对3-账务(是/否)", "核对4-金额(是/否)",
             "核对5-分类(是/否)", "核对6-减值(是/否)", "核对7-利息(是/否)",
             "索引", "是否异常(是/否)", "异常说明",
             "风险等级", "处理建议", "备注"],
            [
                ["2024-12-31", "PZ-001", "计提减值", "资产减值损失",
                 "1503", 100000, 0, "",
                 "减值测算表",
                 "是", "是", "是", "是", "是", "是", "是",
                 "", "否", "", "", "", ""],
            ],
        )
        transport = ASGITransport(app=app)
        async with AsyncClient(transport=transport, base_url="http://test") as client:
            resp = await client.post(
                "/api/workpapers/test-wp/g6-ecl/import-data?sheet=G6-15",
                files={"file": ("test.xlsx", xlsx_bytes, "application/vnd.openxmlformats-officedocument.spreadsheetml.sheet")},
            )
        assert resp.status_code == 200
        data = resp.json()
        payload = data.get("data", data)
        assert payload["imported_count"] > 0

    # --- Invalid sheet ---

    @pytest.mark.asyncio
    async def test_import_data_invalid_sheet(self):
        """Invalid sheet code → 400."""
        xlsx_bytes = _make_xlsx_bytes(["col1"], [["val"]])
        transport = ASGITransport(app=app)
        async with AsyncClient(transport=transport, base_url="http://test") as client:
            resp = await client.post(
                "/api/workpapers/test-wp/g6-ecl/import-data?sheet=G6-99",
                files={"file": ("test.xlsx", xlsx_bytes, "application/vnd.openxmlformats-officedocument.spreadsheetml.sheet")},
            )
        assert resp.status_code == 400


# ═══════════════════════════════════════════════════════════════════════════════
# 3. AI接口: 5个section各返回合理response structure
# ═══════════════════════════════════════════════════════════════════════════════


class TestAiEndpoints:
    """验证 5 个 AI section 端点返回正确结构."""

    @pytest.mark.asyncio
    async def test_ai_stage_conclusion(self, monkeypatch):
        async def _fake_chat(**_kwargs):
            return "三阶段划分审计结论初稿"

        monkeypatch.setattr(
            "app.routers.wp_render_strategies._g6_other_bond_investment_ecl_ai.chat_completion",
            _fake_chat,
        )
        transport = ASGITransport(app=app)
        async with AsyncClient(transport=transport, base_url="http://test") as client:
            resp = await client.post(
                "/api/workpapers/test-wp/g6-ecl/ai/stage-conclusion",
                json={"existingContent": "", "relatedContext": {}},
            )
        assert resp.status_code == 200
        payload = resp.json().get("data", resp.json())
        assert payload["content"] == "三阶段划分审计结论初稿"

    @pytest.mark.asyncio
    async def test_ai_impairment_conclusion(self, monkeypatch):
        async def _fake_chat(**_kwargs):
            return "减值准备测算审计结论初稿"

        monkeypatch.setattr(
            "app.routers.wp_render_strategies._g6_other_bond_investment_ecl_ai.chat_completion",
            _fake_chat,
        )
        transport = ASGITransport(app=app)
        async with AsyncClient(transport=transport, base_url="http://test") as client:
            resp = await client.post(
                "/api/workpapers/test-wp/g6-ecl/ai/impairment-conclusion",
                json={"existingContent": "", "relatedContext": {}},
            )
        assert resp.status_code == 200
        payload = resp.json().get("data", resp.json())
        assert payload["content"] == "减值准备测算审计结论初稿"

    @pytest.mark.asyncio
    async def test_ai_ecl_measurement_conclusion(self, monkeypatch):
        async def _fake_chat(**_kwargs):
            return "ECL计量测试审计结论初稿"

        monkeypatch.setattr(
            "app.routers.wp_render_strategies._g6_other_bond_investment_ecl_ai.chat_completion",
            _fake_chat,
        )
        transport = ASGITransport(app=app)
        async with AsyncClient(transport=transport, base_url="http://test") as client:
            resp = await client.post(
                "/api/workpapers/test-wp/g6-ecl/ai/ecl-measurement-conclusion",
                json={"existingContent": "", "relatedContext": {}},
            )
        assert resp.status_code == 200
        payload = resp.json().get("data", resp.json())
        assert payload["content"] == "ECL计量测试审计结论初稿"

    @pytest.mark.asyncio
    async def test_ai_reversal_writeoff_conclusion(self, monkeypatch):
        async def _fake_chat(**_kwargs):
            return "转回核销审计结论初稿"

        monkeypatch.setattr(
            "app.routers.wp_render_strategies._g6_other_bond_investment_ecl_ai.chat_completion",
            _fake_chat,
        )
        transport = ASGITransport(app=app)
        async with AsyncClient(transport=transport, base_url="http://test") as client:
            resp = await client.post(
                "/api/workpapers/test-wp/g6-ecl/ai/reversal-writeoff-conclusion",
                json={"existingContent": "", "relatedContext": {}},
            )
        assert resp.status_code == 200
        payload = resp.json().get("data", resp.json())
        assert payload["content"] == "转回核销审计结论初稿"

    @pytest.mark.asyncio
    async def test_ai_voucher_conclusion(self, monkeypatch):
        async def _fake_chat(**_kwargs):
            return "凭证检查审计结论初稿"

        monkeypatch.setattr(
            "app.routers.wp_render_strategies._g6_other_bond_investment_ecl_ai.chat_completion",
            _fake_chat,
        )
        transport = ASGITransport(app=app)
        async with AsyncClient(transport=transport, base_url="http://test") as client:
            resp = await client.post(
                "/api/workpapers/test-wp/g6-ecl/ai/voucher-conclusion",
                json={"existingContent": "", "relatedContext": {}},
            )
        assert resp.status_code == 200
        payload = resp.json().get("data", resp.json())
        assert payload["content"] == "凭证检查审计结论初稿"

    @pytest.mark.asyncio
    async def test_ai_invalid_section(self):
        """Invalid section name → 400."""
        transport = ASGITransport(app=app)
        async with AsyncClient(transport=transport, base_url="http://test") as client:
            resp = await client.post(
                "/api/workpapers/test-wp/g6-ecl/ai/invalid-section",
                json={"existingContent": "", "relatedContext": {}},
            )
        assert resp.status_code == 400

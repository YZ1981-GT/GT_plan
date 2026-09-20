"""D4-29/30/31/32 IPO/舞弊组 导入导出 Round-Trip 与结构守卫

Feature: d4-ipo-fraud-writeback-formula-io
Property 1（Requirements 1.1-1.5）：四表 export/import 往返后稳定 id、客户粒度、
全部字段、问卷多选和分组归属不丢失；未知映射不被猜测。

行为级测试（非纯字符串）：直接驱动 _d4_import_export 的 parser/reshape/exporter 纯函数，
构造结构化数据 → 导出 xlsx → 重新解析 → 断言结构等价。用 hypothesis max_examples=5（项目约定）。
"""

from __future__ import annotations

import io

import openpyxl
from hypothesis import given, settings
from hypothesis import strategies as st

from app.routers.wp_render_strategies import _d4_import_export as mod

# ═══════════════════════════════════════════════════════════════════════════════
# 辅助：把「导出行构造」逻辑抽出复用（与 d4_export_data 循环体同源的最小复刻）
# 真正的端点是 async + DB，这里只验证 parser/reshape/exporter 纯函数的结构保真。
# ═══════════════════════════════════════════════════════════════════════════════


def _headers(sheet: str) -> list[str]:
    return mod._get_headers(sheet)


# ─── D4-29 客户信息：一客户一行往返 ───────────────────────────────────────────

_name_st = st.text(min_size=1, max_size=12, alphabet=st.characters(categories=("L", "N")))
_val_st = st.text(max_size=20, alphabet=st.characters(categories=("L", "N", "P", "Zs")))


@settings(max_examples=5)
@given(
    customers=st.lists(
        st.fixed_dictionaries({
            "name": _name_st,
            "creditCode": _val_st,
            "isRelated": st.sampled_from(["是", "否", "待确认", ""]),
            "isAlsoSupplier": st.sampled_from(["是", "否", ""]),
            "infoSource": _val_st,
        }),
        min_size=1, max_size=6,
    ),
)
def test_d4_29_customer_round_trip(customers: list[dict]) -> None:
    """**Validates: Requirements 1.1, 1.2**

    D4-29 客户粒度：导出一客户一行 → 解析回 {id,name,fields}，客户数、name、
    映射字段值不丢失（每客户 = 独立行，非产品汇总冒充）。
    """
    headers = _headers("D4-29")
    # 构造导出行（复刻端点 elif sheet == "D4-29" 分支）
    ws_rows = []
    for c in customers:
        fields = {k: v for k, v in c.items() if k != "name"}
        row = [c["name"]] + [fields.get(mod._D4_29_HEADER_TO_KEY[h], "") for h in headers if h in mod._D4_29_HEADER_TO_KEY]
        ws_rows.append(row)

    # 往返：写 xlsx → 读回 → parser
    wb = openpyxl.Workbook(); ws = wb.active; ws.append(headers)
    for r in ws_rows:
        ws.append(r)
    buf = io.BytesIO(); wb.save(buf); buf.seek(0)
    wb2 = openpyxl.load_workbook(buf, data_only=True); ws2 = wb2.active
    actual_headers = [str(c.value).strip() if c.value else "" for c in next(ws2.iter_rows(min_row=1, max_row=1))]

    parsed = []
    for row in ws2.iter_rows(min_row=2, values_only=True):
        d = mod._parse_d4_29_row(row, actual_headers)
        if d:
            parsed.append(d)

    # 文本经 _safe_str(strip) 归一化：断言归一化后等价（客户粒度/字段不丢）
    def _norm(s: str) -> str:
        return mod._safe_str(s)

    assert len(parsed) == len(customers)  # 客户粒度不丢
    for orig, got in zip(customers, parsed):
        assert got["name"] == _norm(orig["name"])
        assert "id" in got and got["id"].startswith("cust-")  # 稳定 id
        assert got["fields"].get("creditCode", "") == _norm(orig["creditCode"])
        assert got["fields"].get("isRelated", "") == _norm(orig["isRelated"])
        assert got["fields"].get("isAlsoSupplier", "") == _norm(orig["isAlsoSupplier"])


# ─── D4-30 访谈汇总：转置 + 自定义维度往返 ─────────────────────────────────────


@settings(max_examples=5)
@given(
    customers=st.lists(_name_st, min_size=1, max_size=4, unique=True),
    custom_dims=st.lists(st.text(min_size=1, max_size=8, alphabet=st.characters(categories=("L",))), max_size=3, unique=True),
)
def test_d4_30_transpose_and_custom_dims_round_trip(customers: list[str], custom_dims: list[str]) -> None:
    """**Validates: Requirements 1.3**

    D4-30 转置矩阵不被压平：客户名 + 全部固定维度 + 自定义维度往返保留。
    """
    headers = list(_headers("D4-30")) + custom_dims
    # 构造结构化数据
    structured = {
        "customers": [
            {"id": f"iv-{i}", "name": nm, "fields": {"time": f"2026-0{i%9+1}", "reason": "第X大客户", **{cd: f"{nm}-{cd}" for cd in custom_dims}}}
            for i, nm in enumerate(customers)
        ],
        "customDimensions": [{"key": cd, "label": cd} for cd in custom_dims],
    }
    # 导出（复刻端点 elif sheet == "D4-30"）
    wb = openpyxl.Workbook(); ws = wb.active; ws.append(headers)
    for c in structured["customers"]:
        f = c["fields"]
        rv = [c["name"]]
        for h in headers:
            if h == "客户名称":
                continue
            if h in mod._D4_30_HEADER_TO_KEY:
                rv.append(f.get(mod._D4_30_HEADER_TO_KEY[h], ""))
            else:
                rv.append(f.get(h, ""))
        ws.append(rv)
    buf = io.BytesIO(); wb.save(buf); buf.seek(0)
    wb2 = openpyxl.load_workbook(buf, data_only=True); ws2 = wb2.active
    actual_headers = [str(c.value).strip() if c.value else "" for c in next(ws2.iter_rows(min_row=1, max_row=1))]

    rows = []
    for row in ws2.iter_rows(min_row=2, values_only=True):
        d = mod._parse_d4_30_row(row, actual_headers)
        if d:
            rows.append(d)
    reshaped = mod._reshape_d4_30_rows(rows)

    assert len(reshaped["customers"]) == len(customers)  # 客户维度不丢
    got_names = [c["name"] for c in reshaped["customers"]]
    assert got_names == customers
    # 自定义维度恢复（key=label，不猜测归并到固定维度）
    got_dim_keys = {d["key"] for d in reshaped["customDimensions"]}
    assert got_dim_keys == set(custom_dims)
    for orig_nm, c in zip(customers, reshaped["customers"]):
        assert c["fields"].get("time", "").startswith("2026")
        for cd in custom_dims:
            assert c["fields"].get(cd, "") == f"{orig_nm}-{cd}"


# ─── D4-31 问卷：q1_relation 多选数组 + 单对象往返 ─────────────────────────────

_relation_opts = ["客户是终端客户", "客户是经销商(客户)", "客户是供应商", "客户既是客户又是供应商", "客户是ABC的关联方"]


@settings(max_examples=5)
@given(
    q1=st.lists(st.sampled_from(_relation_opts), max_size=5, unique=True),
    q3a=st.sampled_from(["是", "否", ""]),
    q5=st.text(max_size=30),
)
def test_d4_31_questionnaire_multiselect_round_trip(q1: list[str], q3a: str, q5: str) -> None:
    """**Validates: Requirements 1.4**

    D4-31 问卷单对象 + q1_relation 始终 string[]（导出 ; 连接 → 导入拆回数组）。
    """
    data = {"target": "XYZ公司", "q1_relation": q1, "q3a_hasContract": q3a, "q5_otherMatters": q5}
    resp = mod._export_d4_31_questionnaire("D4-31", data)
    # StreamingResponse body_iterator → bytes
    import asyncio

    async def _collect() -> bytes:
        chunks = []
        async for c in resp.body_iterator:
            chunks.append(c if isinstance(c, bytes) else c.encode())
        return b"".join(chunks)

    raw = asyncio.run(_collect())
    wb = openpyxl.load_workbook(io.BytesIO(raw), data_only=True); ws = wb.active
    actual_headers = [str(c.value).strip() if c.value else "" for c in next(ws.iter_rows(min_row=1, max_row=1))]
    assert actual_headers == ["字段", "值"]
    got = mod._parse_d4_31_questionnaire(ws, actual_headers)

    assert isinstance(got.get("q1_relation"), list)  # 始终 string[]，不退化为字符串
    assert got["q1_relation"] == q1
    assert got.get("target") == "XYZ公司"
    assert got.get("q3a_hasContract", "") == q3a
    # 文本字段经 _safe_str(strip) + 非法控制符清洗归一化：断言归一化后等价
    expected_q5 = mod._ILLEGAL_XLSX_CHARS_RE.sub("", q5).strip()
    assert got.get("q5_otherMatters", "") == expected_q5


# ─── D4-32 资金流水：6 组显式分组往返 + 未知组别不归「其他」 ──────────────────


@settings(max_examples=5)
@given(
    counts=st.lists(st.integers(min_value=0, max_value=3), min_size=6, max_size=6),
)
def test_d4_32_group_attribution_round_trip(counts: list[int]) -> None:
    """**Validates: Requirements 1.5**

    D4-32 六组由显式组别列恢复（非按行号硬切）；空组保留、异常/占比原样不重算。
    """
    order = mod._D4_32_GROUP_ORDER
    groups = []
    for gkey, n in zip(order, counts):
        rows = [{"id": f"{gkey}-{i}", "name": f"{gkey}对象{i}", "amount": (i + 1) * 100.0,
                 "ratio": f"{i}%", "bank": "工行", "account": "62xx", "method": "银行流水",
                 "hasAnomaly": "否", "indexRef": ""} for i in range(n)]
        groups.append({"key": gkey, "rows": rows})

    # 导出扁平化（复刻端点 D4-32 分支）
    headers = _headers("D4-32")
    wb = openpyxl.Workbook(); ws = wb.active; ws.append(headers)
    for g in groups:
        glabel = mod._D4_32_GROUP_KEY_TO_LABEL[g["key"]]
        for seq, r in enumerate(g["rows"], start=1):
            ws.append([glabel, seq, r["name"], r["amount"], r["ratio"], r["bank"],
                       r["account"], r["method"], r["hasAnomaly"], r["indexRef"]])
    buf = io.BytesIO(); wb.save(buf); buf.seek(0)
    wb2 = openpyxl.load_workbook(buf, data_only=True); ws2 = wb2.active
    actual_headers = [str(c.value).strip() if c.value else "" for c in next(ws2.iter_rows(min_row=1, max_row=1))]

    parsed = []
    for row in ws2.iter_rows(min_row=2, values_only=True):
        d = mod._parse_d4_32_row(row, actual_headers)
        if d:
            parsed.append(d)
    reshaped = mod._reshape_d4_32_rows(parsed)

    # 六组按固定顺序恢复
    got_keys = [g["key"] for g in reshaped if g["key"] != "__unknown__"]
    assert got_keys == order
    for gkey, n in zip(order, counts):
        g = next(x for x in reshaped if x["key"] == gkey)
        assert len(g["rows"]) == n  # 每组行数按组别列精确恢复，非按行号
        for r in g["rows"]:
            assert r["name"].startswith(gkey)  # 归属正确
            assert r["hasAnomaly"] == "否"      # 人工判断原样保留


@settings(max_examples=5)
@given(unknown_label=st.text(min_size=1, max_size=6, alphabet=st.characters(categories=("L",))))
def test_d4_32_unknown_group_not_auto_other(unknown_label: str) -> None:
    """**Validates: Requirements 1.5, 1.1**

    未知组别不自动归「其他关联方(related)」，落 __unknown__ 供人工映射。
    """
    from hypothesis import assume
    assume(unknown_label not in mod._D4_32_GROUP_LABEL_TO_KEY)

    headers = _headers("D4-32")
    wb = openpyxl.Workbook(); ws = wb.active; ws.append(headers)
    ws.append([unknown_label, 1, "神秘对象", 500.0, "", "", "", "", "是", ""])
    buf = io.BytesIO(); wb.save(buf); buf.seek(0)
    wb2 = openpyxl.load_workbook(buf, data_only=True); ws2 = wb2.active
    actual_headers = [str(c.value).strip() if c.value else "" for c in next(ws2.iter_rows(min_row=1, max_row=1))]

    parsed = [mod._parse_d4_32_row(r, actual_headers) for r in ws2.iter_rows(min_row=2, values_only=True)]
    parsed = [d for d in parsed if d]
    reshaped = mod._reshape_d4_32_rows(parsed)

    related = next(g for g in reshaped if g["key"] == "related")
    assert related["rows"] == []  # 未知组别没被归到「其他」
    unknown = [g for g in reshaped if g["key"] == "__unknown__"]
    assert unknown and len(unknown[0]["rows"]) == 1  # 独立保留待人工映射


# ─── D4-32 未知组保真：label/id/金额三态/账号中文 export→import 往返（Task 9） ──


def _d4_32_flatten_export(parsed: list[dict]) -> list[dict]:
    """复刻端点 elif sheet == "D4-32" 扁平化分支（含未知组保真修复）。

    真源逻辑在 _d4_import_export.d4_export_data 内联，这里最小复刻用于往返验证：
    未知组别行的导出组别列必须用每行保留的原始来源 label（中文），
    不能退化成字面量 "__unknown__"。
    """
    flat: list[dict] = []
    for g in parsed:
        if not isinstance(g, dict):
            continue
        gkey = g.get("key", "")
        glabel = mod._D4_32_GROUP_KEY_TO_LABEL.get(gkey, gkey)
        for seq, r in enumerate(g.get("rows", []), start=1):
            rr = dict(r)
            if gkey == "__unknown__":
                rr["_groupLabel"] = rr.get("groupLabel") or glabel
            else:
                rr["_groupLabel"] = glabel
            rr["_seq"] = seq
            flat.append(rr)
    return flat


# 中文/含账号字符串生成器（含 CJK + 数字 + 常见符号）
_cn_account_st = st.text(
    min_size=1, max_size=16,
    alphabet=st.characters(min_codepoint=0x4E00, max_codepoint=0x9FFF),
) | st.from_regex(r"[0-9\u4e00-\u9fff\-]{1,18}", fullmatch=True)

# 金额三态：0（显式零）/ 空串（空-未知）/ 正数（真实审计金额=2 位小数货币，
# 用 分→元 生成以避免 xlsx 浮点序列化的位级抖动干扰三态语义断言）
_amount_tristate_st = st.sampled_from([0.0, ""]) | st.integers(
    min_value=1, max_value=999_999_999
).map(lambda cents: cents / 100.0)


@settings(max_examples=5)
@given(
    unknown_label=st.text(
        min_size=1, max_size=8,
        alphabet=st.characters(min_codepoint=0x4E00, max_codepoint=0x9FFF),
    ),
    name=st.text(min_size=1, max_size=10, alphabet=st.characters(categories=("L", "N"))),
    amount=_amount_tristate_st,
    account=_cn_account_st,
    bank=st.sampled_from(["中国工商银行", "招商银行", "中国农业银行", ""]),
)
def test_d4_32_unknown_group_fidelity_round_trip(
    unknown_label: str, name: str, amount, account: str, bank: str,
) -> None:
    """**Validates: Requirements 5.2**

    D4-32 未知组保真 export→import 往返：
    - 原始组别 label（中文）保留（不退化成 "__unknown__" 字面量）
    - 稳定 row id 存在且格式正确（ff- 前缀）
    - 金额三态（0 / 空-未知 / 正数）逐态保留，不静默变 0
    - 账号中文/数字文本原样保留
    人工选择已知组前，未知行独立落 __unknown__，不被猜测归并。
    """
    from hypothesis import assume

    # 该 label 必须确实是未知组别（不在六组映射里），且经 _safe_str 后仍非空
    assume(mod._safe_str(unknown_label) not in mod._D4_32_GROUP_LABEL_TO_KEY)
    assume(mod._safe_str(unknown_label) != "")
    assume(mod._safe_str(name) != "")

    # HTML store 结构：__unknown__ 组，行上保留原始来源 label
    store = [
        {"key": k, "rows": []} for k in mod._D4_32_GROUP_ORDER
    ] + [
        {
            "key": "__unknown__",
            "rows": [
                {
                    "id": "ff-seed000001",
                    "name": name,
                    "amount": amount,
                    "ratio": "",
                    "bank": bank,
                    "account": account,
                    "method": "银行流水",
                    "hasAnomaly": "否",
                    "indexRef": "",
                    "groupLabel": unknown_label,  # 原始来源组别（中文）
                }
            ],
        }
    ]

    # 导出扁平化（含保真修复）→ 写 xlsx
    flat = _d4_32_flatten_export(store)
    headers = _headers("D4-32")
    wb = openpyxl.Workbook(); ws = wb.active; ws.append(headers)
    for dr in flat:
        _amt = dr.get("amount", "")
        ws.append([
            dr.get("_groupLabel", ""),
            dr.get("_seq", "") or "",
            dr.get("name", ""),
            mod._safe_float(_amt) if _amt not in (None, "") else "",
            dr.get("ratio", ""),
            dr.get("bank", ""),
            dr.get("account", ""),
            dr.get("method", ""),
            dr.get("hasAnomaly", ""),
            dr.get("indexRef", ""),
        ])
    buf = io.BytesIO(); wb.save(buf); buf.seek(0)

    # 读回 → parse → reshape
    wb2 = openpyxl.load_workbook(buf, data_only=True); ws2 = wb2.active
    actual_headers = [str(c.value).strip() if c.value else "" for c in next(ws2.iter_rows(min_row=1, max_row=1))]
    parsed = [mod._parse_d4_32_row(r, actual_headers) for r in ws2.iter_rows(min_row=2, values_only=True)]
    parsed = [d for d in parsed if d]
    reshaped = mod._reshape_d4_32_rows(parsed)

    # 未知组独立保留，未被归入六组任一
    for gkey in mod._D4_32_GROUP_ORDER:
        g = next(x for x in reshaped if x["key"] == gkey)
        assert g["rows"] == [], f"未知组别不应被猜测归入 {gkey}"
    unknown = [g for g in reshaped if g["key"] == "__unknown__"]
    assert unknown and len(unknown[0]["rows"]) == 1
    got = unknown[0]["rows"][0]

    # ① 原始 label（中文）保真，未退化成字面量 "__unknown__"
    assert got["groupLabel"] == mod._safe_str(unknown_label)
    assert got["groupLabel"] != "__unknown__"

    # ② 稳定 row id 存在且格式正确（xlsx 无 id 列，parse 侧生成 ff- 前缀稳定 id）
    assert isinstance(got.get("id"), str) and got["id"].startswith("ff-")

    # ③ 金额三态逐态保留（空↛0、0↛空、正数不被静默清零）
    if amount == "":
        assert got["amount"] == "", "空-未知金额不得静默变 0"
    elif amount == 0.0:
        assert got["amount"] == 0.0  # 显式 0 保留，非退化成空
    else:
        assert got["amount"] != "" and got["amount"] != 0.0, "正数金额不得被静默清零/清空"
        assert abs(float(got["amount"]) - mod._safe_float(amount)) < 0.005  # 2 位货币精度内一致

    # ④ 账号中文/数字文本原样保留
    assert got["account"] == mod._safe_str(account)
    assert got["name"] == mod._safe_str(name)
    assert got["bank"] == mod._safe_str(bank)


# ═══════════════════════════════════════════════════════════════════════════════
# 结构守卫（非往返）：item_id 双侧一致 + 组别真源锁死 + header↔key 覆盖
# ═══════════════════════════════════════════════════════════════════════════════

import re as _re
from pathlib import Path


def _repo_root() -> Path:
    p = Path(__file__).resolve()
    for _ in range(8):
        if (p / "backend").exists() and (p / "audit-platform").exists():
            return p
        p = p.parent
    raise RuntimeError("repo root not found")


def _read(rel: str) -> str:
    return (_repo_root() / rel).read_text(encoding="utf-8")


# 前端组件持久化用的 item_id（真源=各组件 allResponses.set 的第一参数）
_FRONTEND_ITEM_IDS = {
    "D4-29": "D4-29-customers",
    "D4-30": "D4-30-customers",
    "D4-31": "D4-31-interview",
    "D4-32": "D4-32-groups",
}


def test_item_id_double_side_consistency() -> None:
    """**Validates: Requirements 1.1, 4.1**

    后端 import/export item_id 映射必须与前端组件持久化 item_id 逐字一致，
    否则导入写进的 key 前端读不到（假绿高发点）。
    """
    src = _read("backend/app/routers/wp_render_strategies/_d4_import_export.py")
    for sheet, iid in _FRONTEND_ITEM_IDS.items():
        # 后端两处映射（import/export）都必须出现该 item_id
        assert f'item_id = "{iid}"' in src, f"后端缺 {sheet} → {iid} 映射"

    # 前端组件确实用这些 item_id 持久化
    fe = {
        "D4-29": "audit-platform/frontend/src/components/workpaper/composables/useD4CustomerDetail.ts",
        "D4-30": "audit-platform/frontend/src/components/workpaper/d4/ipo/D4TabInterviewSummary.vue",
        "D4-31": "audit-platform/frontend/src/components/workpaper/d4/ipo/D4TabInterviewDetail.vue",
        "D4-32": "audit-platform/frontend/src/components/workpaper/d4/ipo/D4TabFundFlow.vue",
    }
    for sheet, path in fe.items():
        body = _read(path)
        iid = _FRONTEND_ITEM_IDS[sheet]
        assert f"'{iid}'" in body, f"前端 {sheet} 组件未用 item_id {iid}"


def test_d4_32_group_labels_locked_to_frontend() -> None:
    """**Validates: Requirements 1.5**

    后端 D4-32 组别 label→key 必须与前端 D4TabFundFlow.GROUPS 六组逐字一致
    （双向锁死，防组别错位/未知组别被误归）。
    """
    from app.routers.wp_render_strategies import _d4_import_export as mod

    fe = _read("audit-platform/frontend/src/components/workpaper/d4/ipo/D4TabFundFlow.vue")
    # 前端 GROUPS: { key: 'supplier', label: '主要供应商' } ...
    fe_pairs = dict(_re.findall(r"\{\s*key:\s*'([a-z]+)',\s*label:\s*'([^']+)'", fe))
    for label, key in mod._D4_32_GROUP_LABEL_TO_KEY.items():
        assert fe_pairs.get(key) == label, f"D4-32 组别不一致 backend {key}={label} vs frontend {fe_pairs.get(key)}"
    assert len(mod._D4_32_GROUP_ORDER) == 6


def test_d4_29_header_key_covers_all_fields() -> None:
    """**Validates: Requirements 1.1, 1.2**

    D4-29 表头↔key 映射必须覆盖 _SHEET_HEADERS 中除「客户名称」外的全部列
    （否则导出的列导入时丢字段）。
    """
    from app.routers.wp_render_strategies import _d4_import_export as mod

    headers = [h for h in mod._get_headers("D4-29") if h != "客户名称"]
    for h in headers:
        assert h in mod._D4_29_HEADER_TO_KEY, f"D4-29 表头 {h} 无 key 映射"


def test_d4_31_multi_keys_contains_q1_relation() -> None:
    """**Validates: Requirements 1.4**

    q1_relation 必须在多选集合里（导入必拆回 string[]，不退化字符串）。
    """
    from app.routers.wp_render_strategies import _d4_import_export as mod

    assert "q1_relation" in mod._D4_31_MULTI_KEYS
    # 且 q1_relation 在字段标签映射里
    assert mod._D4_31_LABEL_TO_KEY.get("业务关系(多选)") == "q1_relation"


def test_all_checklist_inserts_include_project_id() -> None:
    """**Validates: Requirements 1.1, 4.1**

    checklist_responses.project_id 是 NOT NULL 列 —— 所有 INSERT 必须带 project_id，
    否则真栈导入必 500（本 spec 真栈实测抓出的既有生产 bug，回归守卫）。
    """
    src = _read("backend/app/routers/wp_render_strategies/_d4_import_export.py")
    inserts = _re.findall(r"INSERT INTO checklist_responses\s*\(([^)]*)\)", src)
    assert inserts, "未找到 INSERT checklist_responses"
    for cols in inserts:
        assert "project_id" in cols, f"INSERT 缺 project_id 列: {cols.strip()}"


def test_d4_32_export_preserves_unknown_group_source_label() -> None:
    """**Validates: Requirements 5.2**

    源级守卫：D4-32 导出扁平化分支对 __unknown__ 组必须用每行保留的原始来源
    label（groupLabel），而非 key→label 查表退化成字面量 "__unknown__"。
    还原该修复即 RED（防往返丢原始中文组别名回归）。
    """
    src = _read("backend/app/routers/wp_render_strategies/_d4_import_export.py")
    # 定位 D4-32 export 扁平化分支
    branch_idx = src.find('elif sheet == "D4-32" and isinstance(parsed, list):')
    assert branch_idx != -1, "未找到 D4-32 导出扁平化分支"
    branch = src[branch_idx:branch_idx + 900]
    # 必须对 __unknown__ 组取每行保留的 groupLabel（保真），不能无条件用查表 label
    assert '__unknown__' in branch, "D4-32 导出分支缺未知组保真处理"
    assert '.get("groupLabel")' in branch, (
        "D4-32 未知组导出必须优先用每行保留的原始来源 label(groupLabel)，"
        "否则往返丢原始中文组别名"
    )

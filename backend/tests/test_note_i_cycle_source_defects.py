"""源模板缺陷登记（Requirement 10）+ md5 冻结 + 生成脚本常量形态守卫。

从 `test_note_i_cycle_structure.py` 拆出（该文件 1671 行超 pre-commit 的 800 行门禁）。
本组判据的共同特征：**对「源模板的已知缺陷」与「生成脚本的常量」断言**，
而不是对模板 JSON 的产出断言 —— 后者在核心文件里。
"""
from __future__ import annotations

import re

import openpyxl
import pytest

from tests.test_note_i_cycle_structure import (  # noqa: F401
    FIX,
    _SRC_DIR,
    _SRC_DISCLOSURE_SHEETS,
    _SRC_FILES,
    _TABLE_NAME_MAX_LEN,
    _norm,
    _section,
)


# ═══════════════════════════════════════════════════════════════════════════════
# Requirement 10 / Task 19：源模板缺陷登记的 stale 检测与处置落实
#
# 判据设计要点：
# - **stale 检测**：每条登记的 `source_ref` 指向的单元格必须仍含 `expect_token`
#   （openpyxl 直读）。源模板一改动（缺陷被上游修掉 / 行列漂移）即打红，
#   提醒重新核对而不是继续按旧结论走。
# - **不得反改源 xlsx**（AC 10.7）：守卫只读不写；另有断言钉死「缺陷仍在源模板里」，
#   若有人偷偷改了源模板来「消除」缺陷，本组会红。
# - **处置落实**：`skip` 类必须在 `wp_code_overrides.json` 里真有 `skip` 登记
#   （否则登记表写了 skip 而链路没排除 = 纯注释承诺）。
# ═══════════════════════════════════════════════════════════════════════════════


def _load_defects():
    from app.services.four_table import i_cycle_source_defects as mod

    return mod


def _src_workbook(file_name: str, data_only: bool = False):
    """按登记里的文件名打开源模板。坏公式需 data_only=False 才看得到。"""
    return openpyxl.load_workbook(_SRC_DIR / file_name, data_only=data_only)


def _parse_source_ref(ref: str) -> tuple[str, str, str]:
    """``"<文件>!<sheet>!<单元格>"`` → 三元组。"""
    parts = str(ref).split("!")
    assert len(parts) == 3, f"source_ref 格式应为 `文件!sheet!单元格`，实为 {ref!r}"
    return parts[0], parts[1], parts[2]


def test_source_defects_registry_shape():
    """登记表结构自检：6 条、id 唯一、循环合法、处置合法、字段非空。"""
    mod = _load_defects()
    defects = mod.I_CYCLE_SOURCE_DEFECTS
    assert len(defects) == 6, (
        f"Requirement 10 立项了 6 处源模板缺陷（AC 10.1~10.6），实为 {len(defects)} 条"
    )
    ids = [d.defect_id for d in defects]
    assert len(set(ids)) == len(ids), f"defect_id 重复：{ids}"
    for d in defects:
        assert d.cycle in _SRC_FILES, f"{d.defect_id} 的 cycle={d.cycle!r} 非 I1~I6"
        assert d.disposition in mod.DISPOSITION_LABELS, (
            f"{d.defect_id} 的 disposition={d.disposition!r} 不在 {list(mod.DISPOSITION_LABELS)}"
        )
        assert d.title.strip(), f"{d.defect_id} 缺 title"
        assert d.evidence.strip(), f"{d.defect_id} 缺 evidence（实测证据）"
        assert d.rationale.strip(), f"{d.defect_id} 缺 rationale（处置理由）"
        assert d.source_ref, f"{d.defect_id} 缺 source_ref"


def test_source_defects_cover_all_acceptance_criteria():
    """六条 AC 逐条有对应登记（按 defect_id 点名，防漏登或改名后失联）。"""
    mod = _load_defects()
    expected = {
        "i1_adjudication_tab_missing_suffix",       # AC 10.1
        "i2_listed_ref_error_continuation_table",   # AC 10.2
        "i1_listed_disposal_typo_in_formula",       # AC 10.3
        "i6_index_cross_workbook_by_design",        # AC 10.4
        "i5_index_sequence_number_missing",         # AC 10.5
        "i3_hidden_sheet_stale_market_return",      # AC 10.6
    }
    got = {d.defect_id for d in mod.I_CYCLE_SOURCE_DEFECTS}
    assert got == expected, f"登记项与 AC 不一一对应：缺 {expected - got}，多 {got - expected}"


def test_source_defects_source_ref_resolvable():
    """反向自检：每条 `source_ref` 的文件与 sheet 必须真实存在（否则 stale 检测在空转）。"""
    mod = _load_defects()
    checked = 0
    for d in mod.I_CYCLE_SOURCE_DEFECTS:
        for ref in d.source_ref:
            file_name, sheet, cell = _parse_source_ref(ref)
            path = _SRC_DIR / file_name
            assert path.exists(), f"{d.defect_id}: 源文件不存在 {file_name}"
            wb = _src_workbook(file_name)
            assert sheet in wb.sheetnames, (
                f"{d.defect_id}: sheet {sheet!r} 不在 {file_name} 里（现有 {wb.sheetnames}）"
            )
            # 单元格可寻址（越界会抛）
            wb[sheet][cell]
            checked += 1
    assert checked >= len(mod.I_CYCLE_SOURCE_DEFECTS), (
        f"只校验了 {checked} 个 source_ref ⇒ 判据在空转"
    )


def test_source_defects_expect_token_not_stale():
    """stale 检测：带 `expect_token` 的登记，其单元格必须仍含该文本。

    🔴 打红不代表代码有 bug —— 而是**源模板变了**（缺陷被上游修掉 / 行列漂移），
    需重新核对该条登记的处置是否还成立，而不是删掉断言。
    """
    mod = _load_defects()
    offenders: list[str] = []
    checked = 0
    for d in mod.I_CYCLE_SOURCE_DEFECTS:
        if not d.expect_token:
            continue
        hit = False
        for ref in d.source_ref:
            file_name, sheet, cell = _parse_source_ref(ref)
            val = str(_src_workbook(file_name)[sheet][cell].value or "")
            checked += 1
            if d.expect_token in val:
                hit = True
        if not hit:
            offenders.append(f"{d.defect_id}: 期望 {d.expect_token!r} 未在 {list(d.source_ref)} 命中")
    assert checked >= 3, f"带 expect_token 的登记只校验了 {checked} 格 ⇒ 判据在空转"
    assert not offenders, f"source_ref stale：{offenders}"


def test_i1_adjudication_tab_name_missing_suffix_still_holds():
    """AC 10.1：I1 的审定表 tab 名仍缺 `-1`，其余五个循环仍带 `-1`。"""
    got: dict[str, list[str]] = {}
    for cycle, path in _SRC_FILES.items():
        got[cycle] = [s for s in openpyxl.load_workbook(path).sheetnames if "审定表" in s]
    assert got["I1"] == ["审定表I1"], f"I1 审定表 tab 名已变：{got['I1']}"
    for cycle in ("I2", "I3", "I4", "I5", "I6"):
        assert got[cycle] == [f"审定表{cycle}-1"], f"{cycle} 审定表 tab 名已变：{got[cycle]}"


def test_i2_listed_ref_error_cell_count_frozen():
    """AC 10.2：I2 上市披露续表区的 `#REF!` 格数冻结为 **20**（非 spec 写的 15）。"""
    wb = _src_workbook("I2 开发支出.xlsx")
    sheet = next(s for s in wb.sheetnames if "披露" in s and "上市" in s)
    ws = wb[sheet]
    coords = [
        c.coordinate
        for row in ws.iter_rows()
        for c in row
        if isinstance(c.value, str) and "#REF!" in c.value
    ]
    assert len(coords) == 20, (
        f"`#REF!` 实测 {len(coords)} 格（冻结基线 20）：{coords}。"
        f"🔴 spec AC 10.2 与 tasks.md 写「15 格」与源模板不符，以本条实测为准"
    )
    rows = sorted({int(re.sub(r"[A-Z]", "", x)) for x in coords})
    cols = sorted({re.sub(r"[0-9]", "", x) for x in coords})
    assert rows == [35, 36, 37, 38, 39], f"行范围漂移：{rows}"
    assert cols == ["A", "B", "C", "D"], f"列范围漂移：{cols}"
    # 续表语义锚点（按意图实现的依据）
    assert _norm(ws["A33"].value) == "续："
    assert _norm(ws["A34"].value) == "项目"
    assert _norm(ws["A40"].value) == "合计"


def test_i1_listed_disposal_typo_scope_frozen():
    """AC 10.3：I1 上市「（1）处置」行的笔误范围冻结 —— B/C 正确、D~L 共 9 格误写「购置」。

    同时钉死 row14（「（1）购置」）的 11 格用「购置」是**正确的**，
    防后续会话把它一起「修正」掉。
    """
    wb = _src_workbook("I1 无形资产、累计摊销及减值准备.xlsx")
    ws = wb[_SRC_DISCLOSURE_SHEETS[("I1", "listed")]]
    assert _norm(ws["A20"].value) == "（1）处置", f"A20 语义已变：{ws['A20'].value!r}"
    assert _norm(ws["A19"].value) == "3.本期减少金额"
    assert _norm(ws["A14"].value) == "（1）购置", f"A14 语义已变：{ws['A14'].value!r}"
    assert _norm(ws["A13"].value) == "2.本期增加金额"

    def judge_of(row: int, col: int) -> str | None:
        v = ws.cell(row=row, column=col).value
        if not isinstance(v, str):
            return None
        m = re.search(r'="([^"]+)"', v)
        return m.group(1) if m else None

    row20 = {openpyxl.utils.get_column_letter(c): judge_of(20, c) for c in range(2, 13)}
    correct = [k for k, v in row20.items() if v == "处置"]
    typo = [k for k, v in row20.items() if v == "购置"]
    assert correct == ["B", "C"], f"row20 用「处置」的列已变：{correct}"
    assert typo == ["D", "E", "F", "G", "H", "I", "J", "K", "L"], (
        f"row20 误写「购置」的列已变：{typo}（冻结基线 D~L 共 9 格）"
    )
    # row14 全用「购置」是正确的，不得被一起改
    row14 = {openpyxl.utils.get_column_letter(c): judge_of(14, c) for c in range(2, 13)}
    assert all(v == "购置" for v in row14.values()), (
        f"row14（本期增加·购置）的判定值被改动：{row14} —— 那 11 格本来就该是「购置」"
    )


def test_i6_index_cross_workbook_hint_still_present():
    """AC 10.4：I6 目录第 7~14 项仍指向 `I2-4`~`I2-11`，且 E10 的提示原文仍在。"""
    wb = _src_workbook("I6 研发费用.xlsx")
    ws = wb[next(s for s in wb.sheetnames if "目录" in s)]
    idx = [_norm(ws.cell(row=r, column=4).value) for r in range(10, 18)]
    assert idx == [f"I2-{n}" for n in range(4, 12)], (
        f"I6 目录第 7~14 项索引号已变：{idx}（冻结基线 I2-4~I2-11）"
    )
    hint = _norm(ws["E10"].value)
    assert "该部分底稿模板在开发支出底稿中" in hint, (
        f"E10 的跨 workbook 提示原文丢失（实为 {hint[:60]!r}）⇒ "
        f"「有意设计」的证据没了，需重新判定是否仍不该当笔误改"
    )


def test_i5_index_sequence_gap_frozen():
    """AC 10.5：I5 目录 7 行内容 / 6 个序号，缺口恰在 `B4`。"""
    wb = _src_workbook("I5 其他非流动资产.xlsx")
    ws = wb[next(s for s in wb.sheetnames if "目录" in s)]
    content_rows = [
        r for r in range(4, 12)
        if _norm(ws.cell(row=r, column=3).value)
        and _norm(ws.cell(row=r, column=3).value) != "内容"
    ]
    seq_rows = [r for r in content_rows if _norm(ws.cell(row=r, column=2).value)]
    assert len(content_rows) == 7, f"内容行数已变：{content_rows}"
    assert len(seq_rows) == 6, f"有序号的行数已变：{seq_rows}"
    assert set(content_rows) - set(seq_rows) == {4}, (
        f"序号缺口不在 B4：缺 {sorted(set(content_rows) - set(seq_rows))}"
    )
    assert _norm(ws["D4"].value) == "I5A", "B4 缺序号的那行索引号应为 I5A"


def test_i3_hidden_sheet_registered_skip():
    """AC 10.6：I3 的 hidden sheet 仍在，且**已在 `wp_code_overrides.json` 标 skip**。

    🔴 「登记表写了 disposition=skip」不等于链路真的排除了它 ——
    分类链路只认 `_WP_CODE_OVERRIDE.get(sheet_name) == "skip"`（不读 `sheet_state`，
    `GT_Custom` 是 `wp_classification_service` 里的硬编码特例）。故必须两侧都断言。
    """
    mod = _load_defects()
    defect = mod.defect_by_id("i3_hidden_sheet_stale_market_return")
    assert defect is not None and defect.disposition == "skip"

    wb = _src_workbook("I3 商誉.xlsx")
    hidden = [s for s in wb.sheetnames if wb[s].sheet_state != "visible"]
    assert "市场平均收益率2017" in hidden, f"I3 的 hidden sheet 已变：{hidden}"

    from app.services.wp_code_override_loader import load_wp_code_overrides

    overrides = load_wp_code_overrides()
    assert overrides.get("市场平均收益率2017") == "skip", (
        "「市场平均收益率2017」未在 wp_code_overrides.json 标 skip ⇒ "
        "它会被分类成 `H-辅助说明` 底稿 sheet 暴露给审计人员"
        "（库侧实测过：workpaper_sheet_classification 里确有该行）"
    )


def test_source_defects_skip_dispositions_all_wired():
    """所有 `disposition == 'skip'` 的登记都必须在 overrides 里真有 skip（防纯注释承诺）。"""
    from app.services.wp_code_override_loader import load_wp_code_overrides

    mod = _load_defects()
    overrides = load_wp_code_overrides()
    skips = [d for d in mod.I_CYCLE_SOURCE_DEFECTS if d.disposition == "skip"]
    assert skips, "反向自检：无 skip 类登记 ⇒ 本条在空转"
    offenders: list[str] = []
    for d in skips:
        # skip 类的 source_ref 第一段的 sheet 名即待 skip 的 sheet
        _, sheet, _ = _parse_source_ref(d.source_ref[0])
        if overrides.get(sheet) != "skip":
            offenders.append(f"{d.defect_id}: sheet {sheet!r} 在 overrides 里是 {overrides.get(sheet)!r}")
    assert not offenders, f"登记为 skip 但链路未排除：{offenders}"


def test_source_defects_payload_projection():
    """下发投影字段齐备且不外泄 stale 检测细节（`expect_token` / `evidence` 不出现）。"""
    mod = _load_defects()
    payload = mod.defects_payload()
    assert len(payload) == len(mod.I_CYCLE_SOURCE_DEFECTS)
    for item in payload:
        assert set(item) == {
            "defect_id", "cycle", "title", "disposition", "disposition_label", "source_ref",
        }, f"投影字段集漂移：{sorted(item)}"
        assert item["disposition_label"] == mod.DISPOSITION_LABELS[item["disposition"]]


# ─── AC 10.7：源 xlsx 不得被反改（md5 冻结） ──────────────────────────────────


#: I 循环六份源模板的 md5 冻结基线（2026-08-12 实测）。
#:
#: 🔴 这是 AC 10.7「登记 + 守卫钉死，**不得反改源 xlsx**」的最强判据：
#: 前面几条守卫只钉「某个缺陷仍在」，改坏别处照样绿；md5 把整份文件钉死。
#:
#: 打红时的正确处理 = **先确认改动来源**：
#: - 若是致同下发了新版模板 ⇒ 逐条复核 `I_CYCLE_SOURCE_DEFECTS` 的处置是否仍成立，
#:   再更新本基线（连同 evidence 一起更新）
#: - 若是有人为了「修笔误」直接改了 xlsx ⇒ **回滚**（会让平台与事务所模板分叉，
#:   且 `wp_template_init_service` 生成底稿时直接复制它）
_SRC_XLSX_MD5: dict[str, str] = {
    "I1 无形资产、累计摊销及减值准备.xlsx": "329be34a3300d4795fef7ad58c4d43e6",
    "I2 开发支出.xlsx": "c656c764fa2b0ba1e84f050ad4e02985",
    "I3 商誉.xlsx": "937daf9c1f7df82b3d2aba1edc505afc",
    "I4 长期待摊费用.xlsx": "471cc8cd873f42ebbfa4913034762618",
    "I5 其他非流动资产.xlsx": "c88837d82ea2cb8136f92fbbb01cf21a",
    "I6 研发费用.xlsx": "ba37838369729052e4a03bf551a59662",
}


def test_source_xlsx_md5_frozen():
    """AC 10.7：六份源模板逐字节冻结（禁反改源 xlsx）。"""
    import hashlib

    offenders: list[str] = []
    for name, expect in _SRC_XLSX_MD5.items():
        path = _SRC_DIR / name
        assert path.exists(), f"源模板缺失：{name}"
        got = hashlib.md5(path.read_bytes()).hexdigest()
        if got != expect:
            offenders.append(f"{name}: {got} != {expect}")
    assert not offenders, (
        f"源模板被改动：{offenders}\n"
        f"🔴 AC 10.7 明令「登记 + 守卫钉死，不得反改源 xlsx」。"
        f"若是致同下发新版，请逐条复核 I_CYCLE_SOURCE_DEFECTS 的处置后再更新本基线"
    )


def test_source_xlsx_md5_baseline_covers_all_six():
    """反向自检：md5 基线必须覆盖 `_SRC_FILES` 全部六份（漏一份就有缺口）。"""
    assert set(_SRC_XLSX_MD5) == {p.name for p in _SRC_FILES.values()}, (
        f"md5 基线与 _SRC_FILES 不一致：\n"
        f"  基线 = {sorted(_SRC_XLSX_MD5)}\n"
        f"  实际 = {sorted(p.name for p in _SRC_FILES.values())}"
    )


def test_hidden_sheets_skipped_by_full_name():
    """AC 10.6 补强：hidden sheet 一律按**完整 sheet 名**标 skip（按尾码会误杀）。

    🔴 反例：若用 `endswith('2017')` 之类的尾码规则，会连带误杀其它循环里以年份
    结尾的正常 sheet；`GT_Custom` 同理不能按 `startswith('GT')` 匹配。
    """
    from app.services.wp_code_override_loader import load_wp_code_overrides

    overrides = load_wp_code_overrides()
    hidden_names: set[str] = set()
    for name, path in _SRC_FILES.items():
        wb = openpyxl.load_workbook(path)
        for sheet in wb.sheetnames:
            if wb[sheet].sheet_state != "visible":
                hidden_names.add(sheet)
    assert hidden_names == {"GT_Custom", "市场平均收益率2017"}, (
        f"I 循环 hidden sheet 集合已变：{sorted(hidden_names)}"
    )
    for sheet in sorted(hidden_names):
        assert overrides.get(sheet) == "skip", (
            f"hidden sheet {sheet!r} 未按完整名标 skip（实为 {overrides.get(sheet)!r}）"
        )
        # 完整名而非尾码：确认 overrides 里就是这个整串键
        assert sheet in overrides, f"{sheet!r} 不是 overrides 的整串键 ⇒ 可能用了尾码规则"


# ═══════════════════════════════════════════════════════════════════════════════
# Property 15（补强）：对**生成脚本的常量**断言表名形态
#
# 🔴 **2026-08-12 变异检验 P15 判 GREEN 抓出的判据缺陷**：
# 把 `fix_note_i_cycle_structure._I5_TABLE2_NAME` 从 `"合同取得成本"` 改回 90 字符的
# 泄漏名 `_I5_TABLE2_LEAKED_NAME`，**全组守卫无一条打红**。两个原因叠加：
#
# 1. 既有的「表名不得以 `[` 开头 / 长度 > 30」判据读的是**模板 JSON**
#    （`note_template_*.json`）。改脚本常量不重跑脚本 ⇒ JSON 不变 ⇒ 判据看不到。
# 2. `--check` 判据（`test_check_passes`）比对的是「脚本期望 vs JSON 实际」，
#    而两侧**都由同一个 `_I5_TABLE2_NAME` 常量推导** —— 改了常量两边一起变、
#    仍然自洽 ⇒ `--check` 结构上无法发现「脚本的期望值本身错了」。
#
# 这是「自洽性判据」的固有盲区：它只能保证两侧一致，不能保证两侧都对。
# 补法 = 直接对脚本里的**目标表名常量**断言形态（不经 JSON、不经 --check）。
# ═══════════════════════════════════════════════════════════════════════════════


#: 生成脚本里「最终表名」常量的名字 → 该表名必须满足的形态约束
#:
#: 只登记**正名过的**表（源模板里表名被 docx 指引段落污染、脚本负责正名的那些）。
#: 值 = (最大长度, 禁止的前缀元组)
_FIX_TABLE_NAME_CONSTS: dict[str, tuple[int, tuple[str, ...]]] = {
    "_I5_TABLE2_NAME": (_TABLE_NAME_MAX_LEN, ("[", "［")),
}


def test_fix_script_table_name_constants_are_not_paragraphs():
    """生成脚本的目标表名常量必须是**短表名**，不得是段落文本。

    🔴 与「模板 JSON 的表名」那条是**两个独立层级**：这条钉脚本的期望值，
    那条钉数据的实际值。少了这条，改坏脚本常量后 `--check` 因两侧同源而仍自洽。
    """
    mod = FIX
    offenders: list[str] = []
    checked = 0
    for const_name, (max_len, bad_prefixes) in _FIX_TABLE_NAME_CONSTS.items():
        assert hasattr(mod, const_name), (
            f"生成脚本缺常量 {const_name} ⇒ 锚点失效，本判据在空转"
        )
        value = str(getattr(mod, const_name) or "")
        checked += 1
        if not value:
            offenders.append(f"{const_name} 为空")
            continue
        if len(value) > max_len:
            offenders.append(f"{const_name} 长度 {len(value)} > {max_len}：{value[:50]!r}…")
        for prefix in bad_prefixes:
            if value.startswith(prefix):
                offenders.append(f"{const_name} 以 {prefix!r} 开头：{value[:50]!r}…")
        # 段落特征：含句读符号的多半是指引文本而非表名
        if any(ch in value for ch in "，。；：、"):
            offenders.append(f"{const_name} 含句读符号（疑似段落文本）：{value[:50]!r}…")
    assert checked == len(_FIX_TABLE_NAME_CONSTS), f"只校验了 {checked} 个常量"
    assert not offenders, f"生成脚本的表名常量形态异常：{offenders}"


def test_fix_script_leaked_name_only_used_as_alias():
    """泄漏名常量只能当 `aliases`（定位旧表用），**不得**被当成最终表名。

    反向自检：脚本里必须**同时**存在「泄漏名常量」与「正名后的表名常量」，
    且两者不相等 —— 若有人把正名常量直接赋成泄漏名，本条打红。
    """
    mod = FIX
    assert hasattr(mod, "_I5_TABLE2_LEAKED_NAME"), "缺泄漏名常量 ⇒ 无法做本判据"
    leaked = str(mod._I5_TABLE2_LEAKED_NAME)
    final = str(mod._I5_TABLE2_NAME)
    assert len(leaked) > _TABLE_NAME_MAX_LEN, (
        f"泄漏名常量长度 {len(leaked)} 不像段落文本 ⇒ 锚点可能已失效：{leaked[:60]!r}"
    )
    assert final != leaked, (
        "最终表名被赋成了泄漏名 ⇒ 正名失效，任何底稿推送都会产出孤儿子表"
        f"（表名是 sub_table_data 的键）：{final[:60]!r}"
    )
    assert final in leaked or "合同取得成本" in final, (
        f"最终表名与泄漏名语义无关联（{final!r}）⇒ 正名结果疑似写错对象"
    )


def test_fix_script_table_names_match_template_json():
    """脚本常量 ↔ 模板 JSON 双向一致（把两个层级钉在一起）。

    有了这条，改脚本常量而不重跑脚本 ⇒ 与 JSON 不一致 ⇒ 打红；
    重跑了脚本 ⇒ JSON 变成段落名 ⇒ 被既有的「表名形态」判据打红。两条路都堵死。
    """
    mod = FIX
    final = str(mod._I5_TABLE2_NAME)
    for variant in ("listed", "soe"):
        sec = _section("I5", variant)
        names = [t.get("name") for t in sec.get("tables") or []]
        assert final in names, (
            f"脚本常量 _I5_TABLE2_NAME={final!r} 不在模板 JSON 的 I5/{variant} 表名集里"
            f"（实际 {names}）⇒ 脚本与数据脱钩，需重跑 fix_note_i_cycle_structure.py --apply"
        )

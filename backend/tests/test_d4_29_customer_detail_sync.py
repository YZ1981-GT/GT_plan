import io
from pathlib import Path

import pytest
from openpyxl import load_workbook
from app.services.workpaper_sync import phase5_d4_29_customer_detail as d
from app.services.workpaper_sync.contracts import load_contract


@pytest.fixture
def contract():
    return load_contract('d4.revenue_detail')


@pytest.fixture
def template():
    from app.services.workpaper_sync import phase5_d4_revenue_detail as provider
    from app.services.workpaper_sync.excel_instrumentation import instrument_workbook_bytes_multi
    return instrument_workbook_bytes_multi(provider.read_authoritative_template(),
        provider.instrumentation_specs(), gate=provider.excel_carrier_gate()).instrumented_bytes


def customers(count=12):
    return [{'id': f'id{i}', 'name': f'客户{i}', 'fields': {k: f'{k}-{i}' for k in d.FIELD_KEYS}} for i in range(count)]


def test_contract_and_frozen_digest(contract, monkeypatch):
    assert d.assert_mapping_digest()
    expected = d.sheet_payload()
    actual = next(s for s in contract.canonical_payload['sheets'] if s['sheet_key'] == d.SHEET_KEY)
    assert actual['sheet_key'] == expected['sheet_key']
    assert actual['excel_name'] == expected['excel_name']
    assert actual['tables'] == expected['tables']
    from app.services.workpaper_sync import phase5_d4_revenue_detail as provider
    assert contract.instrumentation_definition_sha256 == provider.build_contract_payload()['instrumentation_definition_sha256']
    monkeypatch.setitem(d.FIELD_ROWS, 'chairman', 27)
    with pytest.raises(ValueError, match='drift'):
        d.assert_mapping_digest()


def test_projection_merge_add_delete_rename(contract):
    payload = customers(2)
    payload[0]['name'] = '改名'
    projection = d.build_store_projection(payload, contract=contract)
    projection.assert_matches_contract(contract)
    merged, applied, _, removed = d.merge_projection_into_store(projection=projection, base_payload=customers(3))
    assert merged == payload
    assert removed == {'id2'} and applied > 0
    empty = d.build_store_projection([], contract=contract)
    merged, applied, _, removed = d.merge_projection_into_store(projection=empty, base_payload=payload)
    assert merged == [] and applied > 0 and removed == {'id0', 'id1'}


def test_source_mapping_and_12_customer_roundtrip(template, contract):
    original = load_workbook(io.BytesIO(template))[d.MANAGED_SHEET]
    data = d.materialize_transposed_workbook(template, customers())
    ws = load_workbook(io.BytesIO(data))[d.MANAGED_SHEET]
    assert ws['N10'].value == '客户11'
    assert ws['N9'].value == d.IDENTITY_CARRIER_PREFIX + 'id11'
    assert ws.row_dimensions[9].hidden
    assert ws['C28'].value == 'chairman-0'
    assert ws['C34'].value == 'actualController-0'
    assert ws['C41'].value == 'infoSource-0'
    assert ws.column_dimensions['N'].width == original.column_dimensions['M'].width
    assert ws['N28']._style == ws['M28']._style
    for row in original.iter_rows(min_row=42):
        for cell in row:
            assert ws[cell.coordinate].value == cell.value
            assert ws[cell.coordinate]._style == cell._style
    assert set(map(str, original.merged_cells.ranges)) <= set(map(str, ws.merged_cells.ranges))
    assert d.extract_transposed_workbook(data) == customers()
    projection = d.build_store_projection(d.extract_transposed_workbook(data), contract=contract)
    assert len(projection.row_keys[d.TABLE_KEY]) == 12


def test_delete_rename_and_expand_again(template):
    full = d.materialize_transposed_workbook(template, customers())
    remaining = [customers()[11], customers()[0]]
    remaining[0]['name'] = '同名客户'
    remaining[1]['name'] = '同名客户'
    reduced = d.materialize_transposed_workbook(full, remaining)
    assert d.extract_transposed_workbook(reduced) == remaining
    ws = load_workbook(io.BytesIO(reduced))[d.MANAGED_SHEET]
    assert ws['N9'].value is None and ws['N10'].value is None
    assert d.extract_transposed_workbook(d.materialize_transposed_workbook(reduced, [])) == []
    assert len(d.extract_transposed_workbook(d.materialize_transposed_workbook(reduced, customers(14)))) == 14


def test_dispatcher_files(template, contract, tmp_path):
    artifact = tmp_path / 'workbook.xlsx'
    artifact.write_bytes(template)
    d.materialize_file(artifact, d.build_store_projection(customers(), contract=contract), contract)
    projection = d.extract_file(artifact, contract)
    merged, _, _, _ = d.merge_projection_into_store(projection=projection, base_payload=[])
    assert merged == customers()


def test_invalid_duplicate_identity(template):
    with pytest.raises(ValueError, match='unique'):
        d.materialize_transposed_workbook(template, customers(1) * 2)


def test_production_adapter_dispatch(template, contract, tmp_path, monkeypatch):
    import shutil
    from dataclasses import dataclass
    from types import SimpleNamespace
    from app.services.workpaper_sync.adapters import excel
    from app.services.workpaper_sync.adapters.base import SubstrateRole
    from app.services.workpaper_sync.models import ArtifactKind, ArtifactState

    @dataclass
    class Result:
        artifact_sha256: str = ''
        structure_hash: str = ''

    def generic_materialize(**kwargs):
        shutil.copyfile(kwargs['substrate'], kwargs['output'])
        return SimpleNamespace(result=Result())

    monkeypatch.setattr(excel, 'materialize_projection', generic_materialize)
    empty = d.build_store_projection([], contract=contract)
    monkeypatch.setattr(excel, 'extract_projection', lambda **kw: SimpleNamespace(projection=empty))
    class Host:
        adapter_id = 'd4.revenue_detail'
        baseline_representation = None
        substrate_kind = ArtifactKind.canonical
        substrate_role = SubstrateRole.published_representation
        substrate_state = ArtifactState.published
        capability = None
        definitions = None
        _limits = None
        binding = SimpleNamespace(table_key='primary')
        def _assert_same_contract(self, *a, **kw): pass
        def _all_bindings(self): return (self.binding,)
        def _baseline_pair(self): return None, None
        def _substrate_shape_of(self, artifact):
            return self.substrate_role, self.substrate_kind, self.substrate_state
        # 复用生产合并逻辑：extract 现遍历 resolve_transposed_specs 产多段（D4-29 + D4-12），
        # 不能再返 parts[-1]（会丢 D4-29 段）。合并 values + row_keys。
        _merge_projections = excel.ExcelSyncAdapter._merge_projections
    source, output = tmp_path / 'source.xlsx', tmp_path / 'output.xlsx'
    source.write_bytes(template)
    host = Host()
    excel.ExcelSyncAdapter.materialize(host, substrate=source, output=output,
        projection=d.build_store_projection(customers(), contract=contract), contract=contract)
    projection = excel.ExcelSyncAdapter.extract(host, artifact=output, contract=contract)
    assert len(projection.row_keys[d.TABLE_KEY]) == 12
    assert projection.get(d.stable_key_for('id11', 'name')).value == '客户11'


@pytest.mark.parametrize('row', [10, 11, 28, 41])
def test_missing_identity_with_business_data_is_not_deletion(template, row):
    wb = load_workbook(io.BytesIO(d.materialize_transposed_workbook(template, [])))
    ws = wb[d.MANAGED_SHEET]
    ws.cell(row, 3).value = 'orphaned value'
    out = io.BytesIO()
    wb.save(out)
    with pytest.raises(ValueError, match='missing customer identity'):
        d.extract_transposed_workbook(out.getvalue())


@pytest.mark.parametrize('identity', ['=CONCAT("GT-CUSTOMER-","id0")', 'GT-CUSTOMER-', 'GT-CUSTOMER-a/b'])
def test_invalid_physical_identity_rejected(template, identity):
    wb = load_workbook(io.BytesIO(d.materialize_transposed_workbook(template, customers(1))))
    wb[d.MANAGED_SHEET]['C9'] = identity
    out = io.BytesIO()
    wb.save(out)
    with pytest.raises(ValueError):
        d.extract_transposed_workbook(out.getvalue())


def test_duplicate_physical_identity_rejected(template):
    wb = load_workbook(io.BytesIO(d.materialize_transposed_workbook(template, customers(2))))
    wb[d.MANAGED_SHEET]['D9'] = wb[d.MANAGED_SHEET]['C9'].value
    out = io.BytesIO()
    wb.save(out)
    with pytest.raises(ValueError, match='unique'):
        d.extract_transposed_workbook(out.getvalue())


def _saved(wb):
    out = io.BytesIO()
    wb.save(out)
    return out.getvalue()


@pytest.mark.parametrize('count', [0, 2, 14])
def test_transposed_physical_observation_and_publication(template, contract, count):
    from app.services.workpaper_sync.published_identity_observer import (
        _frozen_sheet_anchors, collect_workbook_structure, recompute_structure_hash,
    )
    from app.services.workpaper_sync.publish_time_structure_hash import compute_structure_hash_from_artifact, anchors_from_instrumentation_specs
    from app.services.workpaper_sync import phase5_d4_revenue_detail as provider
    from app.services.workpaper_sync.contracts import parse_contract
    payload = dict(contract.canonical_payload)
    payload['sheets'] = [s for s in payload['sheets'] if s['sheet_key'] == d.SHEET_KEY]
    focused = parse_contract(payload, adapter_id=contract.contract_id)
    anchors = _frozen_sheet_anchors(provider.instrumentation_definition_payload())
    assert anchors == list(anchors_from_instrumentation_specs(provider.instrumentation_specs()))
    anchors = [a for a in anchors if a['sheet_key'] == d.SHEET_KEY]
    data = d.materialize_transposed_workbook(template, customers(count))
    wb = load_workbook(io.BytesIO(data))
    wb[d.MANAGED_SHEET].title = 'Renamed customer sheet'
    wb.defined_names[d.DEFINED_NAME].attr_text = "'Renamed customer sheet'!" + d.MANAGED_REF
    data = _saved(wb)
    assert d.extract_transposed_workbook(data) == customers(count)
    _, physical, _, structure = collect_workbook_structure(data=data, contract=focused, sheet_anchors=anchors)
    assert physical[d.SHEET_KEY] == 'Renamed customer sheet'
    assert len(structure) == 29
    assert compute_structure_hash_from_artifact(data=data, contract=focused, anchors=anchors) == recompute_structure_hash(contract=focused, observed_structure=structure)
    assert d.extract_transposed_workbook(d.materialize_transposed_workbook(data, [])) == []


@pytest.mark.parametrize('damage', ['missing', 'local', 'duplicate', 'range', 'hidden', 'field'])
def test_transposed_anchor_damage_rejected(template, contract, damage):
    import zipfile
    from xml.etree import ElementTree as ET
    from app.services.workpaper_sync.published_identity_observer import collect_workbook_structure, _frozen_sheet_anchors, ObservedIdentityDriftError
    from app.services.workpaper_sync.contracts import parse_contract, ContractDriftError
    payload = dict(contract.canonical_payload)
    payload['sheets'] = [s for s in payload['sheets'] if s['sheet_key'] == d.SHEET_KEY]
    focused = parse_contract(payload, adapter_id=contract.contract_id)
    data = d.materialize_transposed_workbook(template, [])
    if damage in ('hidden', 'field'):
        wb = load_workbook(io.BytesIO(data))
        if damage == 'hidden':
            wb[d.MANAGED_SHEET].row_dimensions[9].hidden = False
        else:
            del wb[d.MANAGED_SHEET]['C41']
        data = _saved(wb)
    else:
        out = io.BytesIO()
        with zipfile.ZipFile(io.BytesIO(data)) as source, zipfile.ZipFile(out, 'w') as target:
            for info in source.infolist():
                blob = source.read(info.filename)
                if info.filename == 'xl/workbook.xml':
                    root = ET.fromstring(blob)
                    names = root.find('{*}definedNames')
                    node = next(n for n in names if n.get('name') == d.DEFINED_NAME)
                    if damage == 'missing': names.remove(node)
                    if damage == 'local': node.set('localSheetId', '0')
                    if damage == 'duplicate': names.append(ET.fromstring(ET.tostring(node)))
                    if damage == 'range': node.text = node.text.replace('$41', '$40')
                    blob = ET.tostring(root)
                target.writestr(info, blob)
        data = out.getvalue()
    anchors = _frozen_sheet_anchors({'transposed_sheets': [d.sheet_payload()]})
    with pytest.raises((ObservedIdentityDriftError, ContractDriftError)):
        collect_workbook_structure(data=data, contract=focused, sheet_anchors=anchors)


def test_request_and_publisher_share_real_transposed_inventory(template, contract, monkeypatch):
    from types import SimpleNamespace
    from app.services.workpaper_sync import published_identity_observer as obs
    from app.services.workpaper_sync import phase5_d4_revenue_detail as provider
    from app.services.workpaper_sync.publish_time_structure_hash import compute_structure_hash_from_artifact
    from app.services.workpaper_sync.contracts import parse_contract
    payload = dict(contract.canonical_payload)
    payload['sheets'] = [s for s in payload['sheets'] if s['sheet_key'] in (provider.SHEET_KEY, d.SHEET_KEY)]
    focused = parse_contract(payload, adapter_id=contract.contract_id)
    frozen = provider.instrumentation_definition_payload()
    frozen['managed_sheets'] = frozen['managed_sheets'][:1]
    data = d.materialize_transposed_workbook(template, [])
    host = SimpleNamespace(_identity_context=lambda **kwargs: {})
    facts = obs.PublishedIdentityObserver._observe_workbook(host, data=data, contract=focused,
        instrumentation=frozen, resolution=None, entry_id=provider.ENTRY_ID, correlation_id='offline')
    assert sum(item[0] == d.SHEET_KEY for item in facts['structure']) == 29
    assert facts['structure_hash'] == compute_structure_hash_from_artifact(data=data, contract=focused, instrumentation=frozen)
    monkeypatch.setattr(obs, 'structure_fingerprint', lambda data: SimpleNamespace(errors=['broken physical collection']))
    with pytest.raises(obs.ArtifactUnreadableError):
        compute_structure_hash_from_artifact(data=data, contract=focused, instrumentation=frozen)
    with pytest.raises(obs.ArtifactUnreadableError):
        obs.PublishedIdentityObserver._observe_workbook(host, data=data, contract=focused,
            instrumentation=frozen, resolution=None, entry_id=provider.ENTRY_ID, correlation_id='offline')


def test_full_contract_publication_hash_includes_d429(template, contract):
    from app.services.workpaper_sync import phase5_d4_revenue_detail as provider
    from app.services.workpaper_sync.publish_time_structure_hash import compute_structure_hash_from_artifact
    data = d.materialize_transposed_workbook(template, [])
    assert len(compute_structure_hash_from_artifact(data=data, contract=contract,
        instrumentation=provider.instrumentation_definition_payload())) == 64

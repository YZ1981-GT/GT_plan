"""其他模块以用户身份改附注时盖「最近一次编辑」（spec disclosure-note-template-word-realignment Task 16~18，design §5）。

每个写入方一条行为测试：真 SQLite + 真 ORM，能走路由的走真路由（鉴权依赖真跑）。
集团模板下发是裸 SQL（SQLite 下 uuid 形态与 ORM 不一致、也不拒绝把字母写进整数列）⇒ 见
``test_note_group_template_stamp_pg.py``。上年导入 / 继承见 ``test_note_prior_year_import_stamp.py``。
"""
from __future__ import annotations

import uuid
from collections.abc import AsyncGenerator
from datetime import datetime, timedelta, timezone
from unittest.mock import patch

import pytest
import pytest_asyncio
import sqlalchemy as sa
from fastapi import FastAPI
from httpx import ASGITransport, AsyncClient
from sqlalchemy.ext.asyncio import AsyncSession, async_sessionmaker, create_async_engine

import tests.conftest  # noqa: F401  注册全部模型 + SQLite JSONB 兼容
import tests.procedure_trim.test_note_linkage_single_source as _linkage
from app.core.database import get_db
from app.deps import get_current_user
from app.models.base import Base, PermissionLevel, ProjectStatus, ProjectUserRole
from app.models.audit_log_models import AuditLogEntry
from app.models.core import Project, ProjectUser, User
from app.models.report_models import DisclosureNote, GroupNoteTemplateBaseline, NoteSectionVersionTree
from app.models.staff_models import ProjectAssignment, StaffMember
from app.routers import group_note_baseline, note_offline, note_wp_mapping, recycle_bin
from tests.test_note_text_batch_endpoints import _Actor

stub_env = _linkage.stub_env  # 程序裁剪联动的替身 fixture（pytest 按属性名注册）

YEAR = 2025
OLD = datetime(2020, 1, 1, tzinfo=timezone.utc)
PREV = uuid.UUID("00000000-0000-0000-0000-0000000000aa")  # 上一位编辑人


def _utc(value: datetime) -> datetime:
    return value.replace(tzinfo=timezone.utc) if value.tzinfo is None else value.astimezone(timezone.utc)


def _cash_table(values: list) -> dict:
    """多表章节：``_tables`` 是真源，顶层 ``rows`` 是第一张表的镜像。"""
    table = {"name": "货币资金", "headers": ["项目", "期末余额", "期初余额"],
             "rows": [{"label": "库存现金", "values": list(values)}, {"label": "合计", "values": [None, None]}]}
    mirror = {"headers": table["headers"],
              "rows": [{"label": "库存现金", "values": list(values)}, {"label": "合计", "values": [None, None]}]}
    return {"_tables": [table], **mirror}


def _edited_offline_package(env: dict, values: list) -> bytes:
    """以 DB 当前值签发 1.1 包，再只改可编辑的 worksheet 值。"""
    from io import BytesIO

    from openpyxl import load_workbook

    from app.services.note_offline_export_service import export_sections_to_xlsx

    payload, _ = export_sections_to_xlsx(
        [{
            "section_id": "八、1", "note_section": "八、1", "section_title": "货币资金",
            "guidance_text": "旧指引", "table_data": _cash_table([100, 80]),
        }],
        project_name="测试", year=YEAR,
        package_context={"project_id": env["ids"]["p"], "year": YEAR, "template_type": "soe"},
    )
    wb = load_workbook(BytesIO(payload))
    ws = next(sheet for sheet in wb.worksheets if sheet.title not in {"注意事项", "章节清单", "_meta_"})
    row = next(r for r in range(1, ws.max_row + 1) if ws.cell(r, 1).value == "库存现金")
    for index, value in enumerate(values, start=2):
        ws.cell(row, index, value=value)
    output = BytesIO()
    wb.save(output)
    return output.getvalue()


@pytest_asyncio.fixture
async def env() -> AsyncGenerator[dict, None]:
    engine = create_async_engine("sqlite+aiosqlite:///:memory:")
    async with engine.begin() as conn:
        await conn.run_sync(lambda c: Base.metadata.create_all(c, tables=[
            User.__table__, Project.__table__, ProjectUser.__table__, StaffMember.__table__,
            ProjectAssignment.__table__, DisclosureNote.__table__, GroupNoteTemplateBaseline.__table__,
            NoteSectionVersionTree.__table__, AuditLogEntry.__table__,
        ]))
    factory = async_sessionmaker(engine, class_=AsyncSession, expire_on_commit=False)
    ids = {k: uuid.uuid4() for k in ("me", "viewer", "p", "q", "cash", "other", "gone", "linked", "baseline")}
    async with factory() as s:
        for key, name in (("p", "甲_2025"), ("q", "乙_2025")):
            s.add(Project(id=ids[key], name=name, client_name=name[:1] + "公司", template_type="soe",
                          report_scope="standalone", status=ProjectStatus.execution))
        await s.flush()
        stamp = {"updated_at": OLD, "updated_by": PREV}
        s.add_all([
            ProjectUser(project_id=ids["p"], user_id=ids["me"], role=ProjectUserRole.auditor,
                        permission_level=PermissionLevel.edit),
            ProjectUser(project_id=ids["p"], user_id=ids["viewer"], role=ProjectUserRole.auditor),
            DisclosureNote(id=ids["cash"], project_id=ids["p"], year=YEAR, note_section="八、1", section_title="货币资金",
                           table_data=_cash_table([100, 80]), guidance_text="旧指引", **stamp),
            DisclosureNote(id=ids["other"], project_id=ids["q"], year=YEAR, note_section="八、1",
                           section_title="货币资金", table_data=_cash_table([1, 2]), **stamp),
            DisclosureNote(id=ids["gone"], project_id=ids["p"], year=YEAR, note_section="八、9", section_title="其他应收款",
                           is_deleted=True, **stamp),
            DisclosureNote(id=ids["linked"], project_id=ids["p"], year=YEAR, note_section="八、2", section_id="soe_bill",
                           section_title="应收票据", text_content="<p>旧正文</p>", **stamp),
        ])
        await s.flush()
        s.add(GroupNoteTemplateBaseline(id=ids["baseline"], name="集团基线", parent_project_id=ids["q"],
                                        template_type="soe", sections_data=[
            {"section_id": "soe_bill", "section_title": "应收票据", "text_content": "<p>基线正文</p>"},
            {"section_id": "soe_new", "note_section": "八、3", "section_title": "应收账款", "text_content": "<p>新</p>"},
            {"section_id": "soe_gone", "note_section": "八、9", "section_title": "其他应收款"},  # 编号被软删行占着
        ]))
        await s.commit()

    app = FastAPI()
    for module in (note_wp_mapping, recycle_bin, note_offline, group_note_baseline):
        app.include_router(module.router)

    async def _override_db():
        async with factory() as session:
            yield session

    app.dependency_overrides[get_db] = _override_db
    actor = {"user": _Actor(ids["me"])}
    app.dependency_overrides[get_current_user] = lambda: actor["user"]
    async with AsyncClient(transport=ASGITransport(app=app), base_url="http://test") as client:
        yield {"client": client, "ids": ids, "factory": factory, "actor": actor}
    await engine.dispose()


async def _note(env: dict, key: str) -> DisclosureNote:
    async with env["factory"]() as db:
        return await db.get(DisclosureNote, env["ids"][key])


def _assert_stamped(note: DisclosureNote, who: uuid.UUID, since: datetime) -> None:
    assert note.updated_by == who, (note.note_section, note.updated_by)
    assert _utc(note.updated_at) >= since, (note.note_section, note.updated_at)


def _assert_untouched(note: DisclosureNote) -> None:
    assert (note.updated_by, _utc(note.updated_at)) == (PREV, OLD), (note.note_section, note.updated_by)


def _since() -> datetime:
    return datetime.now(timezone.utc) - timedelta(seconds=1)


# ── 单元格模式切换（note_wp_mapping.toggle_mode）────────────────────────────────


def _toggle_url(env: dict, note: str = "cash", project: str = "p") -> str:
    return f"/api/disclosure-notes/{env['ids'][project]}/{YEAR}/{env['ids'][note]}/toggle-mode"


class TestToggleCellMode:
    @pytest.mark.asyncio
    async def test_manual_value_lands_in_sidecar_and_stamps(self, env):
        since = _since()
        resp = await env["client"].post(_toggle_url(env), json={
            "row_label": "库存现金", "col_index": 1, "mode": "manual", "manual_value": "88.5"})
        assert resp.status_code == 200, resp.text
        assert resp.json()["changed"] is True
        cash = await _note(env, "cash")  # 新会话读回：真落库（原来原地改 JSONB、从不落库）
        for rows in (cash.table_data["_tables"][0]["rows"], cash.table_data["rows"]):  # 真源 + 顶层镜像
            assert rows[0]["values"] == [100, 88.5], "人工值要写回 values（原来换成字典塞进 values）"
            assert rows[0]["_cell_modes"] == {"1": "manual"}
            assert rows[1] == {"label": "合计", "values": [None, None]}
        _assert_stamped(cash, env["ids"]["me"], since)

    @pytest.mark.asyncio
    async def test_unchanged_mode_is_not_an_edit(self, env):
        resp = await env["client"].post(_toggle_url(env), json={"row_label": "库存现金", "col_index": 0, "mode": "auto"})
        assert resp.status_code == 200 and resp.json()["changed"] is False  # 缺键即自动
        _assert_untouched(await _note(env, "cash"))

    @pytest.mark.asyncio
    async def test_note_of_another_project_is_rejected(self, env):
        """本项目有编辑权，章节 id 却是乙项目的：原来只按 id 取行 ⇒ 改到别的项目。"""
        resp = await env["client"].post(_toggle_url(env, note="other"), json={
            "row_label": "库存现金", "col_index": 1, "mode": "manual", "manual_value": "9"})
        assert resp.status_code == 404
        other = await _note(env, "other")
        assert other.table_data == _cash_table([1, 2]) and other.updated_by == PREV

    @pytest.mark.asyncio
    async def test_readonly_member_and_bad_input_are_rejected(self, env):
        bad = [({"row_label": "不存在的行", "col_index": 0, "mode": "manual"}, 404),
               ({"row_label": "库存现金", "col_index": 2, "mode": "manual"}, 404),
               ({"row_label": "库存现金", "col_index": 0, "mode": "locked"}, 422)]
        for payload, status in bad:
            assert (await env["client"].post(_toggle_url(env), json=payload)).status_code == status, payload
        env["actor"]["user"] = _Actor(env["ids"]["viewer"])
        resp = await env["client"].post(_toggle_url(env), json={"row_label": "库存现金", "col_index": 1, "mode": "manual"})
        assert resp.status_code == 403
        _assert_untouched(await _note(env, "cash"))


# ── 回收站恢复附注（recycle_bin.restore_item）──────────────────────────────────


class TestRecycleBinRestore:
    @pytest.mark.asyncio
    async def test_list_names_note_by_section_code(self, env):
        resp = await env["client"].get("/api/recycle-bin", params={"item_type": "disclosure_note"})
        assert resp.status_code == 200, resp.text
        assert [i["name"] for i in resp.json()["items"]] == ["八、9"], "原来取不存在的 section_code ⇒ 只显示 id 前 8 位"

    @pytest.mark.asyncio
    async def test_restore_stamps_restorer(self, env):
        since = _since()
        resp = await env["client"].post(f"/api/recycle-bin/disclosure_note/{env['ids']['gone']}/restore")
        assert resp.status_code == 200, resp.text
        gone = await _note(env, "gone")
        assert gone.is_deleted is False
        _assert_stamped(gone, env["ids"]["me"], since)


# ── 集团基线下发（group_note_baseline.apply_group_baseline）────────────────────


class TestGroupBaselineApply:
    @pytest.mark.asyncio
    async def test_apply_stamps_updated_and_created_notes_and_skips_occupied_codes(self, env):
        since = _since()
        resp = await env["client"].post(f"/api/projects/{env['ids']['p']}/apply-group-baseline",
                                        json={"baseline_id": str(env["ids"]["baseline"]), "year": YEAR})
        assert resp.status_code == 200, resp.text
        body = resp.json()
        assert body["success"] is True and body["skipped_occupied"] == ["八、9"]
        linked = await _note(env, "linked")
        assert linked.text_content == "<p>基线正文</p>"
        _assert_stamped(linked, env["ids"]["me"], since)
        async with env["factory"]() as db:
            created = (await db.execute(sa.select(DisclosureNote).where(
                DisclosureNote.project_id == env["ids"]["p"], DisclosureNote.year == YEAR,
                DisclosureNote.note_section == "八、3"))).scalar_one()
            same_code = (await db.execute(sa.select(sa.func.count()).select_from(DisclosureNote).where(
                DisclosureNote.project_id == env["ids"]["p"], DisclosureNote.note_section == "八、9"))).scalar_one()
        _assert_stamped(created, env["ids"]["me"], since)
        assert same_code == 1, "被软删行占着的编号又插了一行（唯一索引含软删 ⇒ 真库整次撞键失败）"
        gone = await _note(env, "gone")
        assert gone.is_deleted is True
        _assert_untouched(gone)


# ── 离线导入回写编制指引（note_offline.execute_import）──────────────────────────


class TestOfflineImportGuidance:
    @pytest.mark.asyncio
    async def test_execute_commits_guidance_and_stamps_only_changed_rows(self, env):
        import json

        from app.services import note_offline_import_service as svc

        validation = svc.ImportValidationResult(
            valid=True, meta_data={"八、1": {"guidance_text": "新指引"}, "八、2": {"guidance_text": None}},
            manifest={"template_type": "soe"}, legacy_unsigned=False,
        )
        since = _since()
        with patch.object(svc, "validate_import_file", return_value=validation), \
                patch.object(svc, "context_errors", return_value=[]), \
                patch.object(svc, "apply_import", return_value=svc.ImportResult(success=True)):
            resp = await env["client"].post(
                f"/api/disclosure-notes/{env['ids']['p']}/{YEAR}/offline-import/execute",
                files={"file": ("离线包.xlsx", b"stub", "application/octet-stream")},
                data={"decisions": json.dumps({"八、1": "overwrite", "八、2": "overwrite"})})
        assert resp.status_code == 200, resp.text
        cash = await _note(env, "cash")
        assert cash.guidance_text == "新指引", "路由不提交 ⇒ 会话关闭即回滚，回写从未落库"
        _assert_stamped(cash, env["ids"]["me"], since)
        _assert_untouched(await _note(env, "linked"))  # 指引本来就是空：值没变不算编辑

    @pytest.mark.asyncio
    async def test_execute_writes_tables_reports_actual_changes_and_publishes_event(self, env, monkeypatch):
        """端点真写 DB：回执按实际变化计数，提交后逐章发 NOTE_SECTION_SAVED；幂等重传报 0。"""
        import json

        from app.models.audit_platform_schemas import EventType
        from app.services import event_bus as event_bus_module

        payload = _edited_offline_package(env, [777, 666])
        published = []

        async def _capture(self, event):
            published.append(event)

        monkeypatch.setattr(type(event_bus_module.event_bus), "publish", _capture)
        data = {"decisions": json.dumps({"八、1": "overwrite"}), "merge_cells": "{}"}
        files = {"file": ("离线包.xlsx", payload, "application/vnd.openxmlformats-officedocument.spreadsheetml.sheet")}
        since = _since()
        resp = await env["client"].post(
            f"/api/disclosure-notes/{env['ids']['p']}/{YEAR}/offline-import/execute", data=data, files=files,
        )
        assert resp.status_code == 200, resp.text
        assert resp.json()["sections_imported"] == 1
        assert resp.json()["tables_written"] == 1
        assert resp.json()["changed_sections"] == ["八、1"]
        assert len(resp.json()["version_node_ids"]) == 1
        assert resp.json()["audit_log_id"]
        async with env["factory"]() as db:
            version_count = await db.scalar(sa.select(sa.func.count()).select_from(NoteSectionVersionTree))
            audit_count = await db.scalar(sa.select(sa.func.count()).select_from(AuditLogEntry).where(
                AuditLogEntry.action_type == "note_offline_import"))
        assert version_count == audit_count == 1
        cash = await _note(env, "cash")
        assert cash.table_data["_tables"][0]["rows"][0]["values"] == [777, 666]
        assert cash.table_data["rows"] == cash.table_data["_tables"][0]["rows"]
        _assert_stamped(cash, env["ids"]["me"], since)
        assert len(published) == 1
        assert published[0].event_type == EventType.NOTE_SECTION_SAVED
        assert published[0].extra == {"section_code": "八、1", "action": "offline_import"}

        before_stamp = cash.updated_at
        again = await env["client"].post(
            f"/api/disclosure-notes/{env['ids']['p']}/{YEAR}/offline-import/execute", data=data, files=files,
        )
        assert again.status_code == 400
        assert "导出后已变化" in again.text
        assert (await _note(env, "cash")).updated_at == before_stamp
        assert len(published) == 1, "旧基线重传被拒绝，不应发第二次事件"

    @pytest.mark.asyncio
    async def test_execute_rolls_back_everything_when_audit_record_fails(self, env, monkeypatch):
        """审计 / 版本记录不是 best-effort：写不进 ⇒ 400，表格与版本树一行不留，不发事件。"""
        import json

        from app.services import audit_log_helper
        from app.services import event_bus as event_bus_module

        payload = _edited_offline_package(env, [555, 444])
        published = []

        async def _capture(self, event):
            published.append(event)

        async def _audit_down(db, entry):
            raise RuntimeError("哈希链写入失败")

        monkeypatch.setattr(type(event_bus_module.event_bus), "publish", _capture)
        monkeypatch.setattr(audit_log_helper, "append_audit_log", _audit_down)
        before = (await _note(env, "cash")).table_data
        resp = await env["client"].post(
            f"/api/disclosure-notes/{env['ids']['p']}/{YEAR}/offline-import/execute",
            data={"decisions": json.dumps({"八、1": "overwrite"}), "merge_cells": "{}"},
            files={"file": ("离线包.xlsx", payload,
                            "application/vnd.openxmlformats-officedocument.spreadsheetml.sheet")},
        )
        assert resp.status_code == 400, resp.text
        assert "离线导入审计或版本记录失败" in resp.text and "哈希链写入失败" in resp.text
        assert (await _note(env, "cash")).table_data == before
        async with env["factory"]() as db:
            assert await db.scalar(sa.select(sa.func.count()).select_from(NoteSectionVersionTree)) == 0
        assert published == []


# ── 三格式在线表格保存（triple_format_adapter.update_note_from_structure）────────


class TestTripleFormatSave:
    @pytest.mark.asyncio
    async def test_structure_save_stamps_editor_and_skips_recycle_bin(self, env):
        from app.services.triple_format_adapter import DisclosureNoteAdapter

        structure = {"sheets": [{"cells": {"1:0": {"value": "库存现金"}, "1:1": {"value": 5}, "1:2": {"value": 6}}}]}
        since = _since()
        async with env["factory"]() as db:
            for section in ("八、1", "八、9"):
                await DisclosureNoteAdapter.update_note_from_structure(
                    db, env["ids"]["p"], YEAR, section, structure, user_id=env["ids"]["me"])
            await db.commit()
        cash = await _note(env, "cash")
        assert cash.table_data["rows"] == [{"label": "库存现金", "values": [5, 6]}]
        _assert_stamped(cash, env["ids"]["me"], since)
        gone = await _note(env, "gone")
        assert gone.table_data is None
        _assert_untouched(gone)


# ── 程序裁剪联动附注（procedure_trim_note_linkage.apply_note_linkage）──────────

_TRIMMER = "11111111-1111-1111-1111-111111111111"


class TestProcedureTrimLinkage:
    def test_marking_and_revoking_stamp_the_trimmer(self, stub_env):
        m = _linkage._mod()
        since = _since()
        marked = _linkage._Note("五、40")
        out = _linkage._apply(_linkage._FakeSession(scope_rows=_linkage._all_j_trimmed(), notes=[marked]),
                              actor=_TRIMMER)
        assert out["marked"] == 1 and marked.updated_by == _TRIMMER and marked.updated_at >= since
        revoked = _linkage._Note("五、40", is_empty=True, template_lineage={m.LINEAGE_KEY: {"at": "x"}})
        out = _linkage._apply(_linkage._FakeSession(scope_rows=[("J", "J1", "execute")], notes=[revoked]),
                              actor=_TRIMMER)
        assert out["revoked"] == 1 and revoked.updated_by == _TRIMMER and revoked.updated_at >= since

    def test_conflict_section_is_not_stamped(self, stub_env):
        """有内容的章节只报冲突、不动 ⇒ 也不算编辑。"""
        note = _linkage._Note("五、40", text_content="审计师录入的披露正文")
        out = _linkage._apply(_linkage._FakeSession(scope_rows=_linkage._all_j_trimmed(), notes=[note]),
                              actor=_TRIMMER)
        assert out["marked"] == 0 and not hasattr(note, "updated_by")


# ── 公式推送（系统路径：不清空编辑人）────────────────────────────────────────


class TestFormulaPushKeepsEditor:
    @pytest.mark.asyncio
    async def test_event_push_keeps_editor_and_manual_push_records_pusher(self, monkeypatch):
        from tests._formula_push_env import base_entries, make_env, make_sources
        from tests.test_formula_push_engine import _f2_note, _main

        async with make_env(monkeypatch) as fp:
            await fp.seed_entries(base_entries())
            await fp.add_note("五、1", _f2_note())
            async with fp.factory() as db:
                note = (await db.execute(sa.select(DisclosureNote).where(DisclosureNote.project_id == fp.pid))).scalar_one()
                note.updated_by, note.updated_at = PREV, OLD
                await db.commit()
            since = _since()
            await fp.push("TRIAL_BALANCE_UPDATED")  # 事件触发：没有操作人
            note = await fp.note("五、1")
            assert _main(note.table_data)["库存现金"]["end_amount"] == 390, "附注没被推送，判据空转"
            assert note.updated_by == PREV, "事件推送把上一位编辑人清空了"
            assert _utc(note.updated_at) >= since

            pusher = uuid.uuid4()
            fp.sources = make_sources(hall_1001="200")
            await fp.push("manual", triggered_by=pusher)  # 用户点推送
            note = await fp.note("五、1")
            assert _main(note.table_data)["库存现金"]["end_amount"] == 490
            assert note.updated_by == pusher


class TestOfflineImportMergeWire:
    @pytest.mark.asyncio
    async def test_execute_passes_selected_merge_cells_to_multitable_writer(self, env, monkeypatch):
        """HTTP ``merge_cells`` 必须真正传到服务，不得把 MERGE 静默变成零写入。"""
        import json

        payload = _edited_offline_package(env, [777, 666])
        data = {
            "decisions": json.dumps({"八、1": "merge"}),
            "merge_cells": json.dumps({"八、1": ["0:1"]}),
        }
        files = {"file": ("离线包.xlsx", payload, "application/vnd.openxmlformats-officedocument.spreadsheetml.sheet")}
        published = []

        async def _capture(self, event):
            published.append(event)

        from app.services import event_bus as event_bus_module
        monkeypatch.setattr(type(event_bus_module.event_bus), "publish", _capture)
        response = await env["client"].post(
            f"/api/disclosure-notes/{env['ids']['p']}/{YEAR}/offline-import/execute", data=data, files=files,
        )
        assert response.status_code == 200, response.text
        assert response.json()["sections_imported"] == 1 and response.json()["tables_written"] == 1
        row = (await _note(env, "cash")).table_data["_tables"][0]["rows"][0]
        assert row["values"] == [100, 666], "只选的第二个值应更新，未选的首值保留"
        assert len(published) == 1 and published[0].extra["section_code"] == "八、1"

    @pytest.mark.asyncio
    async def test_execute_rejects_empty_merge_selection_without_writing(self, env):
        import json

        payload = _edited_offline_package(env, [777, 666])
        data = {"decisions": json.dumps({"八、1": "merge"}), "merge_cells": "{}"}
        files = {"file": ("离线包.xlsx", payload, "application/vnd.openxmlformats-officedocument.spreadsheetml.sheet")}
        response = await env["client"].post(
            f"/api/disclosure-notes/{env['ids']['p']}/{YEAR}/offline-import/execute", data=data, files=files,
        )
        assert response.status_code == 400
        assert (await _note(env, "cash")).table_data["_tables"][0]["rows"][0]["values"] == [100, 80]

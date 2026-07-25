"""内容控件注入器测试 — deliverable-lineage-content-control Task 2.1.

覆盖 Property 1（Tag↔section_code）、Property 4（fail-open + 可见内容不变）、
Property 5（幂等）。用合成 docx（body 级 ##SECTION## 标记块），不依赖 OnlyOffice。
"""

from __future__ import annotations

from docx import Document
from docx.oxml.ns import qn

from app.services.content_control_injector import wrap_all_sections, wrap_block_range
from app.services.section_anchor_utils import anchor_name
from app.services.word_doc_utils import scan_section_blocks


def _build_doc_with_sections():
    """构造含两节 ##SECTION## 标记块的 docx。"""
    doc = Document()
    doc.add_paragraph("##SECTION:八、1##")
    doc.add_paragraph("货币资金内容 A")
    doc.add_paragraph("货币资金内容 B")
    doc.add_paragraph("##/SECTION:八、1##")
    doc.add_paragraph("##SECTION:五、1##")
    doc.add_paragraph("应收账款内容")
    doc.add_paragraph("##/SECTION:五、1##")
    return doc


def _all_text(doc) -> list[str]:
    # 按文档顺序遍历所有 w:p（含 w:sdt/w:sdtContent 内的），拼接各段 w:t 文本。
    # 注意：python-docx 的 doc.paragraphs 不下钻 SDT，故直接遍历 body 的 w:p 元素。
    texts = []
    for p in doc.element.body.iter(qn("w:p")):
        runs = [t.text or "" for t in p.iter(qn("w:t"))]
        texts.append("".join(runs))
    return texts


def _find_sdt_tags(doc) -> list[str]:
    tags = []
    for sdt in doc.element.body.iter(qn("w:sdt")):
        pr = sdt.find(qn("w:sdtPr"))
        if pr is None:
            continue
        tag_el = pr.find(qn("w:tag"))
        if tag_el is not None:
            tags.append(tag_el.get(qn("w:val")))
    return tags


def test_wrap_all_sections_injects_tagged_sdt():
    doc = _build_doc_with_sections()
    blocks = scan_section_blocks(doc)
    before_text = _all_text(doc)

    count = wrap_all_sections(doc, blocks)

    assert count == 2
    tags = _find_sdt_tags(doc)
    # Property 1：Tag == anchor_name(section_code)
    assert anchor_name("八、1") in tags
    assert anchor_name("五、1") in tags
    # Property 4：可见文本与段落顺序不变
    assert _all_text(doc) == before_text


def test_content_moved_inside_sdtcontent():
    doc = _build_doc_with_sections()
    blocks = scan_section_blocks(doc)
    wrap_all_sections(doc, blocks)

    # 每个 sdt 的 sdtContent 内应含该节段落
    found = False
    for sdt in doc.element.body.iter(qn("w:sdt")):
        content = sdt.find(qn("w:sdtContent"))
        assert content is not None
        texts = [t.text for t in content.iter(qn("w:t"))]
        if any("货币资金内容 A" in (t or "") for t in texts):
            found = True
    assert found


def test_idempotent_no_double_wrap():
    doc = _build_doc_with_sections()
    blocks = scan_section_blocks(doc)
    assert wrap_all_sections(doc, blocks) == 2

    # 重新扫描（元素已移入 sdtContent，body 级不再有 SECTION 块）→ 再次注入应为 0
    blocks2 = scan_section_blocks(doc)
    count2 = wrap_all_sections(doc, blocks2)
    assert count2 == 0
    # sdt 数量仍为 2（未嵌套叠加）
    assert len(_find_sdt_tags(doc)) == 2


def test_wrap_block_range_fail_open_when_not_in_body():
    doc = _build_doc_with_sections()
    body = doc.element.body
    # 传入不在 body 的元素 → 返回 False，不抛
    orphan = doc.add_paragraph("x")._p
    body.remove(orphan)
    assert wrap_block_range(body, orphan, orphan, "sec_x") is False


def test_wrap_block_range_fail_open_reversed_range():
    doc = _build_doc_with_sections()
    blocks = scan_section_blocks(doc)
    body = doc.element.body
    b = blocks[0]
    # open/close 反转 → False
    assert wrap_block_range(body, b.close_el, b.open_el, anchor_name(b.section_code)) is False


def test_wrap_block_range_guards_none():
    doc = _build_doc_with_sections()
    body = doc.element.body
    assert wrap_block_range(body, None, None, "sec_八_1") is False
    assert wrap_block_range(body, body, body, "") is False

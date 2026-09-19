"""B2-12 docx 导出端点测试（Task 11 / B）。

直接调用 export_b2_12_docx 函数（绕过 Depends），验证：
- 返回 docx Response（正确 media_type + RFC5987 文件名）
- 内容含 9 步/7 点表 + 说明（docx 可被 python-docx 重新打开解析）
- 结构化数据缺失时不报错（空表单导出）
"""
from __future__ import annotations

import io
import uuid
from unittest.mock import AsyncMock, MagicMock

import pytest
from docx import Document

from app.routers.wp_editor_router import export_b2_12_docx


def _make_db(*, checklist_rows=None, proj=("北京测试科技有限公司", 2025)):
    db = AsyncMock()

    async def _execute(stmt, params=None):
        sql = str(getattr(stmt, "text", stmt)).lower()
        if "checklist_responses" in sql:
            res = MagicMock()
            res.fetchall.return_value = checklist_rows or []
            return res
        if "from projects" in sql:
            res = MagicMock()
            res.first.return_value = proj
            return res
        res = MagicMock()
        res.fetchall.return_value = []
        res.first.return_value = None
        return res

    db.execute = _execute
    return db


@pytest.mark.asyncio
async def test_export_docx_with_data():
    rows = [
        ("B2-12-steps", '[{"record":"李某，执业10年"}]'),
        ("B2-12-conclusions", '[{"exists":"是","measure":"扩大程序","impact":"增加期初审计范围"}]'),
        ("B2-12-note", "综合评价说明：前任独立性存在威胁，已应对。"),
    ]
    db = _make_db(checklist_rows=rows)
    resp = await export_b2_12_docx(uuid.uuid4(), uuid.uuid4(), db=db, current_user=MagicMock())

    assert resp.media_type == "application/vnd.openxmlformats-officedocument.wordprocessingml.document"
    assert "filename*" in resp.headers["Content-Disposition"]
    # docx 可被重新解析，含关键内容
    doc = Document(io.BytesIO(resp.body))
    text = "\n".join(p.text for p in doc.paragraphs)
    table_text = "\n".join(
        c.text for t in doc.tables for r in t.rows for c in r.cells
    )
    assert "对前任注册会计师的评价底稿" in text
    assert "北京测试科技有限公司" in text
    assert "李某，执业10年" in table_text
    assert "扩大程序" in table_text
    assert "综合评价说明" in text


@pytest.mark.asyncio
async def test_export_docx_empty():
    db = _make_db(checklist_rows=[], proj=("某公司", 2024))
    resp = await export_b2_12_docx(uuid.uuid4(), uuid.uuid4(), db=db, current_user=MagicMock())
    assert resp.status_code == 200
    doc = Document(io.BytesIO(resp.body))
    # 9 步 + 7 点表头仍在（1 抬头行 + 9 数据行 = 10；7 点同理）
    assert len(doc.tables) == 2
    assert len(doc.tables[0].rows) == 10  # 表头 + 9 步
    assert len(doc.tables[1].rows) == 8   # 表头 + 7 点

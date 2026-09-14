"""G7 长期股权投资(权益法组) — API 集成测试（导入导出 + AI）.

Validates: Requirements 7.3, 7.4

测试清单:
1. G7-4 export-template → source-aligned single data sheet + validations
2. G7-13/G7-14/G7-16 export-template → 200 + multi-sheet xlsx (2 sheets)
3. G7-5/G7-15/G7-17 export-template → 200 + single-sheet xlsx
4. AI invalid section → 400
"""

from __future__ import annotations

import io
from unittest.mock import AsyncMock, MagicMock

import pytest
from httpx import ASGITransport, AsyncClient
from openpyxl import load_workbook

from app.deps import get_current_user
from app.main import app
from app.models.base import UserRole
from app.core.database import get_db


class _FakeUser:
    id = "test-user-id"
    name = "Test User"
    email = "test@example.com"
    role = UserRole.admin


@pytest.fixture(autouse=True)
def override_deps():
    """Mock DB and auth dependencies for all tests."""
    mock_db = AsyncMock()
    mock_db.commit = AsyncMock()
    mock_db.execute = AsyncMock(return_value=MagicMock(fetchall=lambda: [], fetchone=lambda: None))
    mock_db.flush = AsyncMock()

    async def _override_db():
        yield mock_db

    app.dependency_overrides[get_current_user] = lambda: _FakeUser()
    app.dependency_overrides[get_db] = _override_db
    yield mock_db
    app.dependency_overrides.clear()


# ═══════════════════════════════════════════════════════════════════════════════
# 1. G7-4原底稿单sheet导出
# ═══════════════════════════════════════════════════════════════════════════════

@pytest.mark.asyncio
async def test_export_g7_4_source_aligned_template():
    transport = ASGITransport(app=app)
    async with AsyncClient(transport=transport, base_url="http://test") as client:
        resp = await client.post(
            "/api/workpapers/test-wp/g7-equity-method/export-template?sheet=G7-4"
        )

    assert resp.status_code == 200
    wb = load_workbook(io.BytesIO(resp.content))
    assert "G7-4" in wb.sheetnames
    ws = wb["G7-4"]
    headers = [cell.value for cell in ws[2]]
    assert headers[:6] == [
        "投资关系",
        "公司名称",
        "被投资单位ID",
        "级次（国企适用）",
        "企业类型（国企适用）",
        "是否为本期新纳入合并范围的子公司（国企适用）",
    ]
    validations = {
        str(validation.sqref): validation.formula1
        for validation in ws.data_validations.dataValidation
    }
    assert validations["E3:E502"] == '"1,2,3,4,5"'
    assert validations["F3:F502"] == '"是,否"'
    wb.close()


@pytest.mark.asyncio
async def test_g7_4_normalize_preserves_id_and_stamps_ratio_scale():
    from app.routers.wp_render_strategies._g7_long_term_equity_method_import_export import (
        _normalize_g7_4_rows,
    )

    rows, errors = _normalize_g7_4_rows([{
        "id": "keep-me",
        "groupType": "合营企业（共同控制）",
        "investeeName": "甲合营",
        "directHoldingRatio": 40,
        "votingRatio": 40,
    }])
    assert errors == []
    assert rows[0]["id"] == "keep-me"
    assert rows[0]["groupType"] == "joint_venture"
    assert rows[0]["ratioScale"] == "percent"
    assert rows[0]["accountingMethod"] == "权益法"


def test_g7_4_merge_ids_by_name_reuses_existing():
    from app.routers.wp_render_strategies._g7_long_term_equity_method_import_export import (
        _merge_g7_4_ids_by_name,
        _normalize_g7_4_rows,
    )

    normalized, errors = _normalize_g7_4_rows([{
        "groupType": "子公司",
        "investeeName": "甲子公司",
        "directHoldingRatio": 80,
    }])
    assert errors == []
    assert "id" not in normalized[0] or not normalized[0].get("id")

    merged, reused = _merge_g7_4_ids_by_name(
        normalized,
        [{"id": "old-甲", "investeeName": "甲子公司"}],
    )
    assert reused == 1
    assert merged[0]["id"] == "old-甲"


# ═══════════════════════════════════════════════════════════════════════════════
# 2. 宽表分sheet导出：G7-13/G7-14/G7-16 → multi-sheet (2 data sheets)
# ═══════════════════════════════════════════════════════════════════════════════

@pytest.mark.asyncio
@pytest.mark.parametrize("sheet_code", ["G7-13", "G7-14", "G7-16"])
async def test_export_template_multi_sheet(sheet_code: str):
    """POST export-template for multi-sheet tables → 200 + xlsx with ≥2 worksheets.

    G7-13(2区段Tab) / G7-14(2区段Tab) / G7-16(2区段Tab)
    导出模板时按区段分sheet，每个表至少2个data worksheets。

    **Validates: Requirements 7.3**
    """
    transport = ASGITransport(app=app)
    async with AsyncClient(transport=transport, base_url="http://test") as client:
        resp = await client.post(
            f"/api/workpapers/test-wp/g7-equity-method/export-template?sheet={sheet_code}"
        )

    assert resp.status_code == 200, f"{sheet_code} failed: {resp.text}"
    assert "spreadsheetml" in resp.headers.get("content-type", "")

    # Parse xlsx and verify multi-sheet structure
    wb = load_workbook(io.BytesIO(resp.content), read_only=True)
    sheet_names = wb.sheetnames
    wb.close()

    # 宽表导出：至少2个data sheet + 可能有编制说明sheet
    # 所以总sheet数 ≥ 2
    assert len(sheet_names) >= 2, (
        f"{sheet_code} expected ≥2 sheets, got {len(sheet_names)}: {sheet_names}"
    )


# ═══════════════════════════════════════════════════════════════════════════════
# 2. 单sheet导出：G7-5/G7-15/G7-17 → single-sheet xlsx
# ═══════════════════════════════════════════════════════════════════════════════

@pytest.mark.asyncio
@pytest.mark.parametrize("sheet_code", ["G7-5", "G7-15", "G7-17"])
async def test_export_template_single_sheet(sheet_code: str):
    """POST export-template for single-sheet tables → 200 + xlsx with 1 data worksheet.

    G7-5(10列) / G7-15(17列，含ID+毛利率+手工覆盖) / G7-17(9列) 不需要区段拆分，导出为单个data sheet。

    **Validates: Requirements 7.3**
    """
    transport = ASGITransport(app=app)
    async with AsyncClient(transport=transport, base_url="http://test") as client:
        resp = await client.post(
            f"/api/workpapers/test-wp/g7-equity-method/export-template?sheet={sheet_code}"
        )

    assert resp.status_code == 200, f"{sheet_code} failed: {resp.text}"
    assert "spreadsheetml" in resp.headers.get("content-type", "")

    # Parse xlsx — single data sheet (may have guidance sheet)
    wb = load_workbook(io.BytesIO(resp.content), read_only=True)
    sheet_names = wb.sheetnames
    wb.close()

    # 单sheet表：主数据sheet名=sheet_code，可能有编制说明
    # 但data sheets只有1个（主sheet）
    # 总数不超过2（主sheet + 可选编制说明）
    assert len(sheet_names) <= 2, (
        f"{sheet_code} expected ≤2 sheets (data + guidance), got {len(sheet_names)}: {sheet_names}"
    )
    # 确认不是多区段拆分格式（不含"+"字样的区段名）
    segment_like = [s for s in sheet_names if "+" in s or "区段" in s]
    assert len(segment_like) == 0, (
        f"{sheet_code} should not have segment-split sheets: {segment_like}"
    )


# ═══════════════════════════════════════════════════════════════════════════════
# 3. AI invalid section → 400
# ═══════════════════════════════════════════════════════════════════════════════

@pytest.mark.asyncio
async def test_ai_invalid_section_returns_400():
    """POST AI with invalid section → 400.

    **Validates: Requirements 7.4**
    """
    transport = ASGITransport(app=app)
    async with AsyncClient(transport=transport, base_url="http://test") as client:
        resp = await client.post(
            "/api/workpapers/test-wp/g7-equity-method/ai/invalid-section",
            json={"existingContent": "", "relatedContext": {}},
        )

    assert resp.status_code == 400


@pytest.mark.asyncio
async def test_ai_unsupported_section_error_message():
    """POST AI with unsupported section → 400 + descriptive error message."""
    transport = ASGITransport(app=app)
    async with AsyncClient(transport=transport, base_url="http://test") as client:
        resp = await client.post(
            "/api/workpapers/test-wp/g7-equity-method/ai/nonexistent-section",
            json={"existingContent": "", "relatedContext": {}},
        )

    assert resp.status_code == 400
    body = resp.json()
    # ResponseWrapperMiddleware may wrap; check detail or message
    detail = body.get("detail", body.get("message", ""))
    assert "不支持" in detail or "nonexistent" in detail.lower() or "section" in detail.lower()


@pytest.mark.asyncio
async def test_export_template_invalid_sheet_returns_400():
    """POST export-template with invalid sheet code → 400."""
    transport = ASGITransport(app=app)
    async with AsyncClient(transport=transport, base_url="http://test") as client:
        resp = await client.post(
            "/api/workpapers/test-wp/g7-equity-method/export-template?sheet=INVALID"
        )

    assert resp.status_code == 400


# ═══════════════════════════════════════════════════════════════════════════════
# 4. G7-17 归一化 / 导入导出公式列
# ═══════════════════════════════════════════════════════════════════════════════

def test_normalize_g7_17_rows_recalculates_recoverable_and_impairment():
    from app.routers.wp_render_strategies._g7_long_term_equity_method_import_export import (
        _normalize_g7_17_rows,
        _prepare_g7_17_rows_export,
    )

    rows = _normalize_g7_17_rows([
        {
            "investeeName": "甲联营",
            "bookValue": 100000,
            "hasImpairmentSign": "是",
            "fvLessDisposalCost": 70000,
            "valueInUse": 82000,
            "recoverableAmount": 1,  # 脏值，应被重算覆盖
            "impairmentAmount": 1,
        },
        {
            "investeeName": "乙联营",
            "bookValue": 50000,
            "hasImpairmentSign": "否",
            "fvLessDisposalCost": 10,
            "valueInUse": 20,
            "recoverableAmount": 99,
            "impairmentAmount": 99,
        },
    ])

    assert rows[0]["hasImpairmentSign"] is True
    assert rows[0]["recoverableAmount"] == 82000.0
    assert rows[0]["impairmentAmount"] == 18000.0

    assert rows[1]["hasImpairmentSign"] is False
    assert rows[1]["recoverableAmount"] == 0.0
    assert rows[1]["fvLessDisposalCost"] == 0.0
    assert rows[1]["valueInUse"] == 0.0
    assert rows[1]["impairmentAmount"] == 0.0

    exported = _prepare_g7_17_rows_export(rows)
    assert exported[0]["hasImpairmentSign"] == "是"
    assert exported[1]["hasImpairmentSign"] == "否"


@pytest.mark.asyncio
async def test_g7_17_import_normalizes_yes_no_and_formulas(monkeypatch, override_deps):
    """导入 G7-17：是/否→bool，可收回/减值按 CAS8 重算后落库。"""
    from openpyxl import Workbook
    import app.routers.wp_render_strategies._g7_long_term_equity_method_import_export as ie_mod

    captured: dict = {}

    async def _fake_upsert(db, wp_id, item_id, rows, field="conclusion", **kwargs):
        captured["wp_id"] = wp_id
        captured["item_id"] = item_id
        captured["field"] = field
        captured["rows"] = rows

    monkeypatch.setattr(ie_mod, "upsert_json_rows", _fake_upsert)

    wb = Workbook()
    ws = wb.active
    ws.title = "G7-17"
    ws.append(["G7-17 减值测试表"])
    ws.append([
        "被投资单位", "账面价值", "可收回金额", "减值迹象",
        "减值金额", "公允价值-处置费用", "使用价值", "审计结论", "索引",
    ])
    ws.append(["测试联营", 100000, 0, "是", 0, 75000, 80000, "需计提", "IDX-1"])
    buf = io.BytesIO()
    wb.save(buf)
    buf.seek(0)

    transport = ASGITransport(app=app)
    async with AsyncClient(transport=transport, base_url="http://test") as client:
        resp = await client.post(
            "/api/workpapers/test-wp/g7-equity-method/import-data?sheet=G7-17",
            files={"file": ("g7-17.xlsx", buf.getvalue(),
                            "application/vnd.openxmlformats-officedocument.spreadsheetml.sheet")},
        )

    assert resp.status_code == 200, resp.text
    body = resp.json()
    data = body.get("data", body)
    assert data.get("ok") is True
    assert data.get("imported_count") == 1
    assert captured["item_id"] == "G7-17-rows"
    row = captured["rows"][0]
    assert row["hasImpairmentSign"] is True
    assert row["recoverableAmount"] == 80000.0
    assert row["impairmentAmount"] == 20000.0
    assert row["investeeName"] == "测试联营"

"""底稿文件路径：生成即落本底稿文件、打开时回写、在线保存不写共享模板（spec procedure-chain-followup-closure R3.3）。

缺陷（2026-10-25 现查真库）：模板集生成直接把 ``gt_template_library.json`` 里的模板库路径写进 ``file_path``
（H1 ⇒ ``wp_templates/H/H1 固定资产.xlsx``，重药安徽 80 行），库里没有的编码（E1 / D2 等多页底稿）留空串（1917 行）；
在线编辑保存按 ``file_path`` 写盘 ⇒ 会把一个项目的编辑写进所有项目共用的模板。打开底稿时 Univer 懒加载落了
``wp_storage/{项目}/{底稿}.xlsx`` 却不记，两边看到的不是同一份文件。

生成与回写判定用 SQLite 真 ORM；打开底稿的路由（原生 SQL 按 str(uuid) 绑定）与在线保存（PG advisory lock）
用真实 PG 回滚事务。底稿存储目录一律指到临时目录；真实模板目录（backend/wp_templates）只读：落文件是复制，
在线保存用临时模板目录（守卫失效时也写不到真实模板）。
"""
from __future__ import annotations

import hashlib
import importlib
import io
import pkgutil
import uuid
from pathlib import Path

import openpyxl
import pytest
import pytest_asyncio
import sqlalchemy as sa
from fastapi import FastAPI
from httpx import ASGITransport, AsyncClient
from sqlalchemy.dialects.sqlite.base import SQLiteTypeCompiler
from sqlalchemy.ext.asyncio import AsyncSession, async_sessionmaker, create_async_engine
from sqlalchemy.pool import StaticPool

SQLiteTypeCompiler.visit_JSONB = SQLiteTypeCompiler.visit_JSON

import app.models  # noqa: E402
from app.core.config import settings  # noqa: E402
from app.core.database import get_db  # noqa: E402
from app.deps import get_current_user  # noqa: E402
from app.models.base import Base, PermissionLevel, ProjectStatus, ProjectUserRole, UserRole  # noqa: E402
from app.models.core import Project, ProjectUser, User  # noqa: E402
from app.models.workpaper_models import WorkingPaper, WpIndex, WpSourceType, WpStatus, WpTemplateSet  # noqa: E402
from app.routers import wp_template, wp_template_files, wp_template_xlsx  # noqa: E402
from app.services import wp_generate_file, wp_template_init_service  # noqa: E402
from app.services.template_engine import TemplateEngine  # noqa: E402
from app.services.workpaper_sync.canonical_paths import TEMPLATE_ROOT, is_within_root  # noqa: E402

for _module in pkgutil.walk_packages(app.models.__path__, "app.models."):
    importlib.import_module(_module.name)  # create_all 需要全部表（外键引用跨模块）

#: 模板库 JSON 给 H1 的是模板目录路径（旧实现原样写进 file_path），E1 没有路径（旧实现留空串）
_H1_TEMPLATE = TEMPLATE_ROOT / "H" / "H1 固定资产.xlsx"
_H1_TEMPLATE_REL = "wp_templates/H/H1 固定资产.xlsx"
_E1_MAIN_SHEET = "货币资金审定表E1-1"


def _sha(path: Path) -> str:
    return hashlib.sha256(path.read_bytes()).hexdigest()


def _sheetnames(path: Path) -> list[str]:
    wb = openpyxl.load_workbook(str(path), read_only=True)
    try:
        return list(wb.sheetnames)
    finally:
        wb.close()


@pytest_asyncio.fixture
async def world(tmp_path, monkeypatch):
    monkeypatch.chdir(tmp_path)
    storage = tmp_path / "wp_storage"
    monkeypatch.setattr(wp_template_init_service, "STORAGE_DIR", storage)
    engine = create_async_engine("sqlite+aiosqlite:///:memory:", poolclass=StaticPool)
    async with engine.begin() as conn:
        await conn.run_sync(Base.metadata.create_all)
    async with async_sessionmaker(engine, class_=AsyncSession, expire_on_commit=False)() as db:
        user = User(id=uuid.uuid4(), username="own_file_user", email="own_file@example.com", hashed_password="x",
                    role=UserRole.auditor, is_active=True)
        pid = uuid.uuid4()
        db.add_all([user, Project(id=pid, name="【测试】底稿文件_2025", client_name="【测试】底稿文件",
                                  status=ProjectStatus.execution, template_type="soe", report_scope="standalone",
                                  audit_year=2025, scenario="normal")])
        await db.flush()
        db.add(ProjectUser(project_id=pid, user_id=user.id, role=ProjectUserRole.auditor,
                           permission_level=PermissionLevel.edit))
        template_set = WpTemplateSet(set_name="【测试】生成落文件", template_codes=["H1", "E1", "ZZ9"])
        db.add(template_set)
        await db.commit()

        api = FastAPI()
        for router in (wp_template.router, wp_template_files.router, wp_template_xlsx.router):
            api.include_router(router)

        async def _db():
            yield db

        async def _user() -> User:
            return user

        api.dependency_overrides[get_db] = _db
        api.dependency_overrides[get_current_user] = _user
        async with AsyncClient(transport=ASGITransport(app=api), base_url="http://own") as http:
            yield {"http": http, "db": db, "pid": pid, "set_id": template_set.id, "storage": storage,
                   "root": tmp_path}
    await engine.dispose()


async def _papers(db: AsyncSession, pid: uuid.UUID) -> dict[str, tuple[uuid.UUID, str]]:
    rows = (await db.execute(
        sa.select(WpIndex.wp_code, WorkingPaper.id, WorkingPaper.file_path)
        .join(WorkingPaper, WorkingPaper.wp_index_id == WpIndex.id)
        .where(WpIndex.project_id == pid)
    )).all()
    return {code: (wid, fp) for code, wid, fp in rows}


async def _add_paper(db: AsyncSession, pid: uuid.UUID, code: str, file_path: str) -> uuid.UUID:
    idx = WpIndex(project_id=pid, wp_code=code, wp_name=f"底稿{code}", audit_cycle=code[0],
                  status=WpStatus.not_started)
    db.add(idx)
    await db.flush()
    wp = WorkingPaper(project_id=pid, wp_index_id=idx.id, file_path=file_path,
                      source_type=WpSourceType.template, file_version=1)
    db.add(wp)
    await db.commit()
    return wp.id


async def _file_path(db: AsyncSession, wp_id: uuid.UUID) -> str:
    return (await db.execute(sa.select(WorkingPaper.file_path).where(WorkingPaper.id == wp_id))).scalar_one()


# ── 模板集生成 ─────────────────────────────────────────────────────────────


async def test_template_set_generation_gives_every_paper_its_own_file(world):
    """H1（模板库给的是模板目录路径）/ E1（模板库没有路径）/ ZZ9（没有模板）都落到本底稿存储位置。"""
    template_sha = _sha(_H1_TEMPLATE)
    r = await world["http"].post(f"/api/projects/{world['pid']}/working-papers/generate",
                                 json={"template_set_id": str(world["set_id"]), "year": 2025})
    assert r.status_code == 200, r.text
    body = r.json()
    assert body["created"] == 3 and body["blank_codes"] == ["ZZ9"]
    assert "没有找到底稿模板" in body["message"] and "ZZ9" in body["message"]

    papers = await _papers(world["db"], world["pid"])
    for code, (wid, file_path) in papers.items():
        expected = world["storage"] / str(world["pid"]) / f"{wid}.xlsx"
        assert file_path == str(expected), f"{code} 的 file_path 应指向本底稿自己的文件：{file_path!r}"
        assert expected.is_file()
        assert not is_within_root(TEMPLATE_ROOT, Path(file_path).resolve())
    assert _sha(Path(papers["H1"][1])) == template_sha, "H1 应是模板的逐字节副本（生成不逐份打开填表头）"
    assert _E1_MAIN_SHEET in _sheetnames(Path(papers["E1"][1])), "E1 应按模板目录索引找到多页模板"
    blank_wb = openpyxl.load_workbook(papers["ZZ9"][1], read_only=True)
    try:
        assert blank_wb.active["A2"].value == "底稿名称: 底稿ZZ9"
        assert blank_wb.active[f"A{wp_generate_file.BLANK_NOTICE_ROW}"].value == wp_generate_file.BLANK_NOTICE
    finally:
        blank_wb.close()
    assert _sha(_H1_TEMPLATE) == template_sha, "生成只能复制模板，不得改动模板目录"


async def test_template_set_with_merged_header_templates_still_generates(world):
    """J2A 的模板首几行有合并单元格：表头写入会抛 AttributeError（约一成模板如此）。

    整批落文件是全有或全无 ⇒ 在请求里逐份填表头会让含这些编码的模板集整体生成失败；生成只复制模板，照常成功。
    """
    db = world["db"]
    template_set = WpTemplateSet(set_name="【测试】合并单元格表头", template_codes=["J2A"])
    db.add(template_set)
    await db.commit()
    r = await world["http"].post(f"/api/projects/{world['pid']}/working-papers/generate",
                                 json={"template_set_id": str(template_set.id), "year": 2025})
    assert r.status_code == 200, r.text
    assert r.json()["created"] == 1 and r.json()["blank_codes"] == []
    wid, file_path = (await _papers(db, world["pid"]))["J2A"]
    assert file_path == str(world["storage"] / str(world["pid"]) / f"{wid}.xlsx")
    assert Path(file_path).is_file()


async def test_generation_failure_discards_files_already_written(world, monkeypatch):
    """同批里某份落文件失败 ⇒ 本批已落的文件全部删掉再抛出（底稿行随事务回滚）。"""
    original = wp_generate_file.materialize_workpaper_file

    def _boom(project_id, wp_id, wp_code, wp_name, **kwargs):
        if wp_code == "E1":
            raise RuntimeError("磁盘写满")
        return original(project_id, wp_id, wp_code, wp_name, **kwargs)

    monkeypatch.setattr(wp_generate_file, "materialize_workpaper_file", _boom)
    with pytest.raises(RuntimeError, match="磁盘写满"):
        await TemplateEngine().generate_project_workpapers_detailed(
            db=world["db"], project_id=world["pid"], template_set_id=world["set_id"], year=2025)
    await world["db"].rollback()
    project_dir = world["storage"] / str(world["pid"])
    assert not project_dir.exists() or list(project_dir.iterdir()) == [], "失败前已落的 H1 文件应被删掉"
    assert await _papers(world["db"], world["pid"]) == {}


# ── 打开底稿时回写 file_path ─────────────────────────────────────────────────


@pytest.mark.parametrize("kind", ["blank", "template", "missing"])
async def test_lazy_bind_when_paper_has_no_own_file(world, kind):
    raw = {"blank": "", "template": _H1_TEMPLATE_REL, "missing": "storage/projects/x/workpapers/H1.xlsx"}[kind]
    wid = await _add_paper(world["db"], world["pid"], "H1", raw)
    target = world["storage"] / "lazy.xlsx"
    assert await wp_generate_file.bind_lazy_workpaper_file(world["db"], wid, target) is True
    await world["db"].commit()
    assert await _file_path(world["db"], wid) == str(target)


async def test_lazy_bind_keeps_existing_own_file(world):
    """``file_path`` 已指向一份存在的本底稿文件（如 legacy storage/projects/...）⇒ 不动，那份文件仍被引用。"""
    legacy = world["root"] / "storage" / "projects" / str(world["pid"]) / "workpapers" / "H" / "H1.xlsx"
    legacy.parent.mkdir(parents=True)
    legacy.write_bytes(_H1_TEMPLATE.read_bytes())
    wid = await _add_paper(world["db"], world["pid"], "H1", str(legacy))
    assert await wp_generate_file.bind_lazy_workpaper_file(world["db"], wid, world["storage"] / "x.xlsx") is False
    assert await _file_path(world["db"], wid) == str(legacy)


async def test_lazy_bind_does_not_overwrite_a_concurrent_write(world):
    """读到原值之后另一次打开先改了 file_path ⇒ 条件更新 0 行，不覆盖对方。"""
    db = world["db"]
    wid = await _add_paper(db, world["pid"], "H1", "")
    other = world["storage"] / "other.xlsx"
    real_execute = db.execute
    calls = {"n": 0}

    async def _racing_execute(statement, *args, **kwargs):
        calls["n"] += 1
        if calls["n"] == 2:  # 本次条件更新之前，另一次打开先写了
            await real_execute(sa.update(WorkingPaper).where(WorkingPaper.id == wid).values(file_path=str(other)))
        return await real_execute(statement, *args, **kwargs)

    db.execute = _racing_execute  # 实例属性遮住类方法；del 后恢复
    try:
        assert await wp_generate_file.bind_lazy_workpaper_file(db, wid, world["storage"] / "mine.xlsx") is False
    finally:
        del db.execute
    assert await _file_path(db, wid) == str(other)


# ── 真实 PG（回滚事务）：路由的原生 SQL 按 str(uuid) 绑定，SQLite 存 32 位十六进制取不到 ───────────


@pytest_asyncio.fixture
async def pg(tmp_path, monkeypatch):
    if not settings.DATABASE_URL.startswith("postgresql"):
        pytest.skip("need PostgreSQL")
    monkeypatch.chdir(tmp_path)  # cwd 是 legacy 根之一：临时模板 / 存储都在它下面
    monkeypatch.setattr(wp_template_init_service, "STORAGE_DIR", tmp_path / "wp_storage")
    engine = create_async_engine(settings.DATABASE_URL, pool_pre_ping=True)
    conn = await engine.connect()
    trans = await conn.begin()
    db = AsyncSession(bind=conn, join_transaction_mode="create_savepoint", expire_on_commit=False)
    try:
        tag = uuid.uuid4().hex[:8]
        user = User(username=f"pcfc_own_{tag}", email=f"pcfc_own_{tag}@example.com", hashed_password="x",
                    role=UserRole.auditor, is_active=True)
        project = Project(name=f"【测试】底稿文件_{tag}", client_name="【测试】底稿文件", status=ProjectStatus.execution,
                          template_type="soe", report_scope="standalone", audit_year=2025, scenario="normal")
        db.add_all([user, project])
        await db.flush()
        db.add(ProjectUser(project_id=project.id, user_id=user.id, role=ProjectUserRole.auditor,
                           permission_level=PermissionLevel.edit))
        await db.flush()

        api = FastAPI()
        for router in (wp_template_files.router, wp_template_xlsx.router):
            api.include_router(router)

        async def _db():
            yield db

        async def _user() -> User:
            return user

        api.dependency_overrides[get_db] = _db
        api.dependency_overrides[get_current_user] = _user
        async with AsyncClient(transport=ASGITransport(app=api), base_url="http://own") as http:
            yield {"http": http, "db": db, "pid": project.id, "user": user, "root": tmp_path,
                   "storage": tmp_path / "wp_storage", "monkeypatch": monkeypatch}
    finally:
        await db.close()
        await trans.rollback()
        await conn.close()
        await engine.dispose()


@pytest.mark.parametrize("route", ["template-file", "template-file/xlsx-to-json"])
async def test_opening_a_paper_records_the_lazily_created_file(pg, route):
    """Univer 打开（取 xlsx / 转 JSON）落了本底稿文件 ⇒ file_path 改指向它（原为空串）。"""
    wid = await _add_paper(pg["db"], pg["pid"], "H1", "")
    r = await pg["http"].get(f"/api/projects/{pg['pid']}/workpapers/{wid}/{route}")
    assert r.status_code == 200, r.text[:300]
    expected = pg["storage"] / str(pg["pid"]) / f"{wid}.xlsx"
    assert expected.is_file()
    assert await _file_path(pg["db"], wid) == str(expected)


# ── 在线保存遇到模板路径：先复制再写 ─────────────────────────────────────────


def _xlsx_bytes(marker: str) -> bytes:
    book = openpyxl.Workbook()
    book.active["A1"] = marker
    buffer = io.BytesIO()
    book.save(buffer)
    return buffer.getvalue()


async def test_online_save_to_template_path_copies_first(pg):
    """底稿 file_path 指向共享模板 ⇒ 保存写进本底稿自己的副本、改写 file_path，模板字节不变。

    模板目录用临时目录（守卫失效时也写不到真实模板）；按生产 provision 入口为 ``wopi_put_file`` lane 发布
    approved bundle（同 test_wopi_put_file_rollback）。
    """
    from app.services import wopi_service
    from app.services.workpaper_sync import canonical_paths
    from app.services.workpaper_sync.artifacts import CanonicalArtifactRepository
    from app.services.workpaper_sync.opaque_entry_gate import authority_model_for_lane
    from app.services.workpaper_sync.repository import WorkpaperSyncRepository
    from app.services.workpaper_sync.resolution import CanonicalResolutionService
    from app.services.workpaper_sync.writer_migration import OpaqueAuthorityProvisioner

    root, db, monkeypatch = pg["root"], pg["db"], pg["monkeypatch"]
    monkeypatch.setattr(canonical_paths, "BACKEND_ROOT", root)  # artifact store 根（调用期取值）
    template_root = root / "wp_templates"
    monkeypatch.setattr(wopi_service, "_TEMPLATE_ROOT", template_root)
    template = template_root / "H" / "H1 固定资产.xlsx"
    template.parent.mkdir(parents=True)
    template.write_bytes(_xlsx_bytes("共享模板"))
    template_bytes = template.read_bytes()

    wid = await _add_paper(db, pg["pid"], "H1", str(template))
    artifacts = CanonicalArtifactRepository(root)
    await OpaqueAuthorityProvisioner(
        session=db, repository=WorkpaperSyncRepository(db), artifacts=artifacts,
        resolution=CanonicalResolutionService(db, artifacts),
    ).provision(project_id=pg["pid"], wp_id=wid, authority_model=authority_model_for_lane("wopi_put_file"))
    await db.commit()

    saved = _xlsx_bytes("项目编辑")
    await wopi_service.WOPIHostService().put_file(db, wid, saved)
    own = pg["storage"] / str(pg["pid"]) / f"{wid}.xlsx"
    assert template.read_bytes() == template_bytes, "共享模板被写了"
    assert own.read_bytes() == saved
    assert await _file_path(db, wid) == str(own)


@pytest.mark.parametrize("method,route", [("GET", ""), ("POST", "/init"), ("GET", "/sheet/Sheet1"),
                                          ("GET", "/xlsx-to-json")])
async def test_lazy_open_rejects_a_workpaper_from_another_project(pg, method, route):
    """项目成员拿另一项目的 wp_id 走本项目 URL ⇒ 四个懒加载入口都 404：不落文件、不回写 file_path。"""
    other = Project(name="【测试】底稿归属隔离", client_name="【测试】底稿归属隔离", status=ProjectStatus.execution,
                    template_type="soe", report_scope="standalone", audit_year=2025, scenario="normal")
    pg["db"].add(other)
    await pg["db"].flush()
    idx = WpIndex(project_id=other.id, wp_code="H1", wp_name="底稿H1", audit_cycle="H", status=WpStatus.not_started)
    pg["db"].add(idx)
    await pg["db"].flush()
    wp = WorkingPaper(project_id=other.id, wp_index_id=idx.id, file_path="", source_type=WpSourceType.template, file_version=1)
    pg["db"].add(wp)
    await pg["db"].flush()

    response = await pg["http"].request(
        method, f"/api/projects/{pg['pid']}/workpapers/{wp.id}/template-file{route}")
    assert response.status_code == 404, response.text[:300]
    assert "不属于当前项目" in response.json()["detail"]
    assert await _file_path(pg["db"], wp.id) == ""
    assert not list(pg["storage"].rglob(f"{wp.id}.*")), "不得为另一项目的底稿落文件"

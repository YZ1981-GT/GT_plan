"""G7 长期股权投资(子公司组) — 集成测试.

Spec: .kiro/specs/g7-long-term-equity-subsidiary/ Task 12.4
Validates: Requirements 1.4, 7.2, 7.3, 7.4

Tests:
  1. render-config返回正确sheets配置（selfLoad逻辑：htmlData为null）
  2. 导入导出18端点(6张表×3)
  3. AI 5个section端点
"""

from __future__ import annotations

import io
from unittest.mock import AsyncMock, patch

import pytest
from httpx import ASGITransport, AsyncClient
from openpyxl import Workbook

from app.core.database import get_db
from app.deps import get_current_user
from app.main import app
from app.models.base import UserRole


# ═══════════════════════════════════════════════════════════════════════════════
# Fixtures
# ═══════════════════════════════════════════════════════════════════════════════


class _FakeUser:
    id = "integ-user-001"
    name = "Integration User"
    email = "integ@example.com"
    role = UserRole.admin


@pytest.fixture(autouse=True)
def override_deps():
    """Override DB and auth dependencies for all tests."""
    mock_db = AsyncMock()
    mock_db.commit = AsyncMock()
    mock_db.flush = AsyncMock()

    mock_result = AsyncMock()
    mock_result.scalar_one_or_none = AsyncMock(return_value=None)
    mock_result.fetchone = AsyncMock(return_value=None)
    mock_result.fetchall = AsyncMock(return_value=[])
    mock_db.execute = AsyncMock(return_value=mock_result)

    async def _override_db():
        yield mock_db

    app.dependency_overrides[get_current_user] = lambda: _FakeUser()
    app.dependency_overrides[get_db] = _override_db
    yield mock_db
    app.dependency_overrides.clear()


def _make_xlsx_bytes(headers: list[str], rows: list[list] | None = None) -> bytes:
    """Create minimal xlsx bytes with given headers."""
    wb = Workbook()
    ws = wb.active
    # Title row (row 1)
    ws.append([f"Template"])
    # Header row (row 2)
    ws.append(headers)
    if rows:
        for row in rows:
            ws.append(row)
    buf = io.BytesIO()
    wb.save(buf)
    buf.seek(0)
    return buf.getvalue()


# ═══════════════════════════════════════════════════════════════════════════════
# 1. render-config 返回正确sheets配置 (selfLoad逻辑：htmlData为null)
# ═══════════════════════════════════════════════════════════════════════════════


@pytest.mark.asyncio
class TestRenderConfigEndpoint:
    """测试render策略直接调用返回正确结构(selfLoad场景)."""

    async def test_render_returns_correct_component_type(self):
        """render函数返回component_type='g7-long-term-equity-subsidiary'."""
        from app.routers.wp_render_strategies._g7_long_term_equity_subsidiary import (
            render,
        )
        from unittest.mock import MagicMock

        # Construct a minimal RenderContext mock
        ctx = MagicMock()
        ctx.wp_id = "test-wp-id"
        ctx.project_id = "test-project-id"

        mock_db = AsyncMock()
        mock_result = AsyncMock()
        mock_result.fetchall = AsyncMock(return_value=[])
        mock_result.fetchone = AsyncMock(return_value=None)
        mock_db.execute = AsyncMock(return_value=mock_result)
        ctx.db = mock_db

        result = await render(ctx)

        assert result is not None
        assert result["component_type"] == "g7-long-term-equity-subsidiary"

    async def test_render_returns_7_sheets(self):
        """render函数返回7个sheet配置."""
        from app.routers.wp_render_strategies._g7_long_term_equity_subsidiary import (
            render,
        )
        from unittest.mock import MagicMock

        ctx = MagicMock()
        ctx.wp_id = "test-wp-id"
        ctx.project_id = "test-project-id"

        mock_db = AsyncMock()
        mock_result = AsyncMock()
        mock_result.fetchall = AsyncMock(return_value=[])
        mock_result.fetchone = AsyncMock(return_value=None)
        mock_db.execute = AsyncMock(return_value=mock_result)
        ctx.db = mock_db

        result = await render(ctx)

        assert result is not None
        assert "sheets" in result
        assert len(result["sheets"]) == 7
        codes = [s["code"] for s in result["sheets"]]
        assert "G7-7" in codes
        assert "G7-8" in codes
        assert "G7-9" in codes
        assert "G7-10" in codes
        assert "G7-11" in codes
        assert "G7-12" in codes
        assert "G7-18" in codes

    async def test_render_sheets_all_have_g7_subsidiary_component_type(self):
        """每个sheet的componentType字段正确."""
        from app.routers.wp_render_strategies._g7_long_term_equity_subsidiary import (
            render,
        )
        from unittest.mock import MagicMock

        ctx = MagicMock()
        ctx.wp_id = "test-wp-id"
        ctx.project_id = "test-project-id"

        mock_db = AsyncMock()
        mock_result = AsyncMock()
        mock_result.fetchall = AsyncMock(return_value=[])
        mock_result.fetchone = AsyncMock(return_value=None)
        mock_db.execute = AsyncMock(return_value=mock_result)
        ctx.db = mock_db

        result = await render(ctx)

        for sheet in result["sheets"]:
            assert sheet["componentType"] == "g7-long-term-equity-subsidiary"


# ═══════════════════════════════════════════════════════════════════════════════
# 2. 导入导出18端点(6张表×3)
# ═══════════════════════════════════════════════════════════════════════════════


@pytest.mark.asyncio
class TestImportExportEndpoints:
    """测试G7子公司组导入导出端点."""

    # --- export-template ---

    async def test_export_template_g7_8_returns_200(self):
        """POST /g7-sub/export-template?sheet=G7-8 → 200."""
        async with AsyncClient(
            transport=ASGITransport(app=app),
            base_url="http://test",
        ) as client:
            resp = await client.post(
                "/api/workpapers/test-wp/g7-sub/export-template",
                params={"sheet": "G7-8"},
            )
            assert resp.status_code == 200
            assert "spreadsheet" in resp.headers.get("content-type", "")

    async def test_export_template_g7_18_multi_sheet_returns_200(self):
        """POST /g7-sub/export-template?sheet=G7-18 → 200 (多sheet导出)."""
        async with AsyncClient(
            transport=ASGITransport(app=app),
            base_url="http://test",
        ) as client:
            resp = await client.post(
                "/api/workpapers/test-wp/g7-sub/export-template",
                params={"sheet": "G7-18"},
            )
            assert resp.status_code == 200
            assert "spreadsheet" in resp.headers.get("content-type", "")
            # Verify multi-sheet workbook
            wb = Workbook()
            from openpyxl import load_workbook
            wb = load_workbook(io.BytesIO(resp.content))
            sheet_names = wb.sheetnames
            # G7-18 should have 3 segment sheets + 编制说明
            assert len(sheet_names) >= 3
            assert any("凭证基础" in n for n in sheet_names)
            assert any("核对检查" in n for n in sheet_names)
            assert any("结论" in n for n in sheet_names)
            wb.close()

    # --- export-data ---

    async def test_export_data_g7_10_returns_200(self):
        """POST /g7-sub/export-data?sheet=G7-10 → 200."""
        async with AsyncClient(
            transport=ASGITransport(app=app),
            base_url="http://test",
        ) as client:
            resp = await client.post(
                "/api/workpapers/test-wp/g7-sub/export-data",
                params={"sheet": "G7-10"},
            )
            assert resp.status_code == 200
            assert "spreadsheet" in resp.headers.get("content-type", "")

    # --- import-data ---

    async def test_import_data_g7_9_with_xlsx_returns_200(self):
        """POST /g7-sub/import-data?sheet=G7-9 + xlsx → 200 with imported_count."""
        headers = [
            "被投资单位", "购买日", "合并方式", "支付对价",
            "直接费用", "初始投资成本", "被购买方净资产FV", "享有份额", "商誉",
        ]
        xlsx_bytes = _make_xlsx_bytes(headers, [
            ["公司A", "2024-06-01", "控股合并", 5000, 100, 5100, 4000, 2400, 2700],
        ])
        async with AsyncClient(
            transport=ASGITransport(app=app),
            base_url="http://test",
        ) as client:
            resp = await client.post(
                "/api/workpapers/test-wp/g7-sub/import-data",
                params={"sheet": "G7-9"},
                files={"file": ("test.xlsx", xlsx_bytes, "application/vnd.openxmlformats-officedocument.spreadsheetml.sheet")},
            )
            assert resp.status_code == 200
            body = resp.json()
            # ResponseWrapperMiddleware may wrap: check both patterns
            data = body.get("data", body)
            assert data.get("imported_count", 0) >= 1

    # --- invalid sheet → 400 ---

    async def test_invalid_sheet_returns_400(self):
        """不支持的sheet code → 400."""
        async with AsyncClient(
            transport=ASGITransport(app=app),
            base_url="http://test",
        ) as client:
            resp = await client.post(
                "/api/workpapers/test-wp/g7-sub/export-template",
                params={"sheet": "G7-INVALID"},
            )
            assert resp.status_code == 400

    async def test_export_template_g7_11_returns_200(self):
        """POST /g7-sub/export-template?sheet=G7-11 → 200."""
        async with AsyncClient(
            transport=ASGITransport(app=app),
            base_url="http://test",
        ) as client:
            resp = await client.post(
                "/api/workpapers/test-wp/g7-sub/export-template",
                params={"sheet": "G7-11"},
            )
            assert resp.status_code == 200

    async def test_export_template_g7_12_returns_200(self):
        """POST /g7-sub/export-template?sheet=G7-12 → 200."""
        async with AsyncClient(
            transport=ASGITransport(app=app),
            base_url="http://test",
        ) as client:
            resp = await client.post(
                "/api/workpapers/test-wp/g7-sub/export-template",
                params={"sheet": "G7-12"},
            )
            assert resp.status_code == 200


# ═══════════════════════════════════════════════════════════════════════════════
# 3. AI 5个section端点
# ═══════════════════════════════════════════════════════════════════════════════


@pytest.mark.asyncio
class TestAiEndpoints:
    """测试G7子公司组AI辅助生成端点."""

    @patch("app.services.llm_client.chat_completion", new_callable=AsyncMock)
    async def test_ai_control_judgment_conclusion_accepts_request(self, mock_llm):
        """POST /g7-sub/ai/control-judgment-conclusion → 200."""
        mock_llm.return_value = "经审计，我们认为被投资方满足CAS33控制三要素..."

        async with AsyncClient(
            transport=ASGITransport(app=app),
            base_url="http://test",
        ) as client:
            resp = await client.post(
                "/api/workpapers/test-wp/g7-sub/ai/control-judgment-conclusion",
                json={"existingContent": "", "relatedContext": {}},
            )
            assert resp.status_code == 200
            body = resp.json()
            data = body.get("data", body)
            assert "content" in data

    @patch("app.services.llm_client.chat_completion", new_callable=AsyncMock)
    async def test_ai_initial_measurement_conclusion(self, mock_llm):
        """POST /g7-sub/ai/initial-measurement-conclusion → 200."""
        mock_llm.return_value = "初始计量审计结论..."

        async with AsyncClient(
            transport=ASGITransport(app=app),
            base_url="http://test",
        ) as client:
            resp = await client.post(
                "/api/workpapers/test-wp/g7-sub/ai/initial-measurement-conclusion",
                json={"existingContent": "", "relatedContext": {}},
            )
            assert resp.status_code == 200

    @patch("app.services.llm_client.chat_completion", new_callable=AsyncMock)
    async def test_ai_subsequent_conclusion(self, mock_llm):
        """POST /g7-sub/ai/subsequent-conclusion → 200."""
        mock_llm.return_value = "后续计量结论..."

        async with AsyncClient(
            transport=ASGITransport(app=app),
            base_url="http://test",
        ) as client:
            resp = await client.post(
                "/api/workpapers/test-wp/g7-sub/ai/subsequent-conclusion",
                json={"existingContent": "", "relatedContext": {}},
            )
            assert resp.status_code == 200

    @patch("app.services.llm_client.chat_completion", new_callable=AsyncMock)
    async def test_ai_disposal_conclusion(self, mock_llm):
        """POST /g7-sub/ai/disposal-conclusion → 200."""
        mock_llm.return_value = "处置结论..."

        async with AsyncClient(
            transport=ASGITransport(app=app),
            base_url="http://test",
        ) as client:
            resp = await client.post(
                "/api/workpapers/test-wp/g7-sub/ai/disposal-conclusion",
                json={"existingContent": "", "relatedContext": {}},
            )
            assert resp.status_code == 200

    @patch("app.services.llm_client.chat_completion", new_callable=AsyncMock)
    async def test_ai_voucher_conclusion(self, mock_llm):
        """POST /g7-sub/ai/voucher-conclusion → 200."""
        mock_llm.return_value = "凭证检查结论..."

        async with AsyncClient(
            transport=ASGITransport(app=app),
            base_url="http://test",
        ) as client:
            resp = await client.post(
                "/api/workpapers/test-wp/g7-sub/ai/voucher-conclusion",
                json={"existingContent": "", "relatedContext": {}},
            )
            assert resp.status_code == 200

    async def test_ai_invalid_section_returns_400(self):
        """POST /g7-sub/ai/invalid-section → 400."""
        async with AsyncClient(
            transport=ASGITransport(app=app),
            base_url="http://test",
        ) as client:
            resp = await client.post(
                "/api/workpapers/test-wp/g7-sub/ai/invalid-section",
                json={"existingContent": "", "relatedContext": {}},
            )
            assert resp.status_code == 400

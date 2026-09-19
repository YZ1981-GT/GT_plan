"""报告正文段落定位器守卫 — deliverable-lineage-wiring-and-writeback-closure Task 17

**Validates: Requirements 9.1, 9.2, 9.6, 12.4**

判据全部以**真实源模板**为准（`backend/data/audit_report_templates/report_body/*.docx`），
不用合成 docx —— 报告正文的章节结构来自致同模板，合成 fixture 只能验证我自己的假设。

两条源模板实证事实（本文件把它们钉死，防后来者按直觉改判据）：

1. **中文序号会随可选段落取舍漂移** —— 「管理层和治理层对财务报表的责任」在
   模板 D 简版是 `三、`、在模板 A 简版是 `五、`。故只能按序号后的名称匹配。
2. **docx 标题不带「段」字** —— `SECTION_ID_MAP` 键是 `审计意见段`，模板写 `一、审计意见`。
"""

from __future__ import annotations

from pathlib import Path

import pytest
from docx import Document

from app.services.report_body_section_blocks import (
    REPORT_BODY_ANCHOR_PREFIX,
    SECTION_NAME_TO_ID,
    build_report_body_sections_payload,
    extract_report_body_sections,
    match_section_heading,
    report_body_anchor_name,
    scan_report_body_sections,
    section_id_from_report_body_anchor,
)
from app.services.report_body_service import SECTION_ID_MAP
from app.services.section_anchor_utils import (
    FOREIGN_ANCHOR_PREFIXES,
    write_section_anchors,
)

TPL_DIR = (
    Path(__file__).resolve().parents[1]
    / "data"
    / "audit_report_templates"
    / "report_body"
)


def _templates(limit: int = 6) -> list[Path]:
    if not TPL_DIR.exists():
        return []
    return sorted(TPL_DIR.glob("*.docx"))[:limit]


requires_templates = pytest.mark.skipif(
    not _templates(), reason="报告正文源模板目录不存在（精简部署）"
)


# ─── 名称表派生（单一真源） ─────────────────────────────────────────────────


def test_name_map_is_derived_from_section_id_map_not_copied():
    """名称表必须由 `SECTION_ID_MAP` 派生 —— 条数一致且值域相同。

    另抄一份会在上游增删章节时静默漂移（改一处另一处不红）。
    """
    assert len(SECTION_NAME_TO_ID) == len(SECTION_ID_MAP)
    assert set(SECTION_NAME_TO_ID.values()) == set(SECTION_ID_MAP.values())
    # 去「段」后的名称确实与原键不同（证明派生规则生效，不是原样拷贝）
    assert "审计意见" in SECTION_NAME_TO_ID
    assert "审计意见段" not in SECTION_NAME_TO_ID
    assert SECTION_NAME_TO_ID["审计意见"] == "opinion"


def test_anchor_namespace_registered_for_isolation():
    """本命名空间必须已在 `section_anchor_utils` 登记为外域前缀。

    不登记 ⇒ 附注回填的 `scan_anchor_blocks` 会把 `sec_rb_*` 反解成伪章节码。
    """
    assert REPORT_BODY_ANCHOR_PREFIX in FOREIGN_ANCHOR_PREFIXES


def test_anchor_name_round_trip():
    for sid in SECTION_ID_MAP.values():
        name = report_body_anchor_name(sid)
        assert name.startswith(REPORT_BODY_ANCHOR_PREFIX)
        assert section_id_from_report_body_anchor(name) == sid
    # 非本域锚点
    assert section_id_from_report_body_anchor("sec_八_1") is None
    assert section_id_from_report_body_anchor("") is None


# ─── 标题匹配：序号漂移 + 不误判正文 ────────────────────────────────────────


@pytest.mark.parametrize(
    ("text", "expected"),
    [
        ("一、审计意见", "opinion"),
        ("二、形成审计意见的基础", "basis"),
        ("三、关键审计事项", "kam"),
        # 同一章节在不同模板变体里序号不同 —— 都必须命中同一 section_id
        ("三、管理层和治理层对财务报表的责任", "mgmt_responsibility"),
        ("五、管理层和治理层对财务报表的责任", "mgmt_responsibility"),
        ("六、注册会计师对财务报表审计的责任", "cpa_responsibility"),
        ("二、形成保留意见的基础", "qualified_basis"),
    ],
)
def test_heading_match_is_ordinal_agnostic(text, expected):
    assert match_section_heading(text) == expected


@pytest.mark.parametrize(
    "text",
    [
        "",
        "我们审计了某公司财务报表。",
        # 正文里提到章节名的句子**不得**被判成标题（模板 A 责任段末尾真有这句）
        "从与治理层沟通过的事项中，我们确定哪些事项因而构成关键审计事项。",
        # 有序号但名称未登记
        "七、我们自拟的一段",
        # 子标题（KAM 内部）不是章节标题
        "（一）X子标题",
    ],
)
def test_heading_match_rejects_non_headings(text):
    assert match_section_heading(text) is None


# ─── 真实模板端到端 ────────────────────────────────────────────────────────


@requires_templates
def test_real_templates_scan_to_known_sections():
    """每个源模板都能扫出 ≥2 个章节，且 section_id 全在已登记值域内。"""
    valid = set(SECTION_ID_MAP.values())
    for tpl in _templates():
        doc = Document(str(tpl))
        secs = scan_report_body_sections(doc)
        assert len(secs) >= 2, f"{tpl.name} 只扫出 {len(secs)} 个章节"
        ids = [s.section_id for s in secs]
        assert set(ids) <= valid, f"{tpl.name} 出现未登记 section_id: {ids}"
        assert len(ids) == len(set(ids)), f"{tpl.name} section_id 重复: {ids}"
        # 每个章节至少有一段正文
        for s in secs:
            assert s.content_els, f"{tpl.name}/{s.section_id} 无正文元素"


@requires_templates
def test_section_text_excludes_heading():
    """章节文字口径**排除标题** —— 标题是模板固定文字且序号会漂移。

    纳入标题会让「换模板变体」被误判成人工编辑。
    """
    for tpl in _templates(3):
        doc = Document(str(tpl))
        for sec in scan_report_body_sections(doc):
            heading = "".join(
                t.text or ""
                for t in sec.heading_el.iter(
                    "{http://schemas.openxmlformats.org/wordprocessingml/2006/main}t"
                )
            ).strip()
            assert heading, "标题段落文字为空（定位器抓错元素）"
            assert heading not in sec.text(), (
                f"{tpl.name}/{sec.section_id}: 章节文字里含标题"
            )


@requires_templates
def test_extract_and_payload_are_consistent():
    """`extract_report_body_sections` 与 `build_report_body_sections_payload` 同源。"""
    for tpl in _templates(3):
        doc = Document(str(tpl))
        flat = extract_report_body_sections(doc)
        payload = build_report_body_sections_payload(doc)

        assert {p["section_id"] for p in payload} == set(flat)
        for p in payload:
            assert p["content"] == flat[p["section_id"]]
            assert set(p) == {
                "section_id",
                "section_name",
                "content",
                "section_order",
            }
            # section_name 必须是带「段」的正式名（对齐 JSON 模式的读取端）
            assert p["section_name"] in SECTION_ID_MAP
        # section_order 连续从 1 起
        assert [p["section_order"] for p in payload] == list(
            range(1, len(payload) + 1)
        )


@requires_templates
def test_anchor_write_is_idempotent_in_section_set():
    """写锚点后再扫，章节集合不变（锚点是不可见元素，不干扰定位）。"""
    for tpl in _templates(3):
        doc = Document(str(tpl))
        before = [s.section_id for s in scan_report_body_sections(doc)]
        write_section_anchors(
            doc,
            [s.to_section_block() for s in scan_report_body_sections(doc)],
            namer=report_body_anchor_name,
        )
        after = [s.section_id for s in scan_report_body_sections(doc)]
        assert before == after, f"{tpl.name}: 写锚点改变了章节定位结果"


# ─── 反向自检 ──────────────────────────────────────────────────────────────


def test_reverse_selfcheck_ordinal_based_id_would_be_wrong():
    """反向自检：若按「序号」当 section 标识，模板 D 与模板 A 会给同一章节不同标识。

    钉死「序号不可作标识」这条源模板事实 —— 一旦有人改成按序号定位，
    跨模板变体的章节状态与回填目标会全部错位。
    """
    d_style = "三、管理层和治理层对财务报表的责任"
    a_style = "五、管理层和治理层对财务报表的责任"
    assert match_section_heading(d_style) == match_section_heading(a_style)
    # 而序号本身确实不同（证明这两个样本有区分度、不是同一字符串）
    assert d_style != a_style


def test_scan_tolerates_docx_without_any_heading():
    """无任何章节标题的 docx ⇒ 返回 `[]`，不抛异常（fail-open）。"""
    doc = Document()
    doc.add_paragraph("封面")
    doc.add_paragraph("目录")
    assert scan_report_body_sections(doc) == []
    assert extract_report_body_sections(doc) == {}
    assert build_report_body_sections_payload(doc) == []

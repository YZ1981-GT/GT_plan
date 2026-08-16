"""test_x3_column_property —— 用例层。

判据（常量 / 登记表 / fixture / 纯函数 / helper）见 `_test_x3_column_property_criteria.py`，
原文件 docstring 也在那里。拆分原因：原单文件 1166 行 > pre-commit 800 行门禁。
"""

from __future__ import annotations

import importlib.util
import io
import json
import re
import string
import sys
import types
from collections.abc import Mapping, Sequence
from dataclasses import dataclass
from functools import lru_cache
from pathlib import Path
from typing import Any, NamedTuple
from unittest import mock

import pytest
from hypothesis import HealthCheck, given, settings
from hypothesis import strategies as st
from openpyxl import load_workbook

from app.routers.wp_render_strategies import _x3_adjustment_import_export as impl

from tests._test_x3_column_property_criteria import (  # noqa: F401  fixtures 需在本模块命名空间
    _AbstractRow,
    _IMPL_NAME,
    _IMPL_PATH,
    _MARK,
    _Product,
    _RUNS_ROUNDTRIP,
    _RUNS_STANDIN,
    _SETTINGS,
    _V_DUP,
    _V_EXTRA,
    _V_MISPLACED,
    _V_MISSING,
    _V_ORDER,
    _V_PLACEHOLDER,
    _V_ROW_WIDTH,
    _V_SHEET_NAME,
    _V_WIDTH,
    _arb_rows,
    _codes,
    _expected_face,
    _extra_keys,
    _extract,
    _header_violations,
    _live,
    _materialize,
    _placeholder_positions,
    _row_violations,
    _spec,
    _template_violations,
)

__all__ = [
    "_AbstractRow",
    "_IMPL_NAME",
    "_IMPL_PATH",
    "_MARK",
    "_Product",
    "_RUNS_ROUNDTRIP",
    "_RUNS_STANDIN",
    "_SETTINGS",
    "_V_DUP",
    "_V_EXTRA",
    "_V_MISPLACED",
    "_V_MISSING",
    "_V_ORDER",
    "_V_PLACEHOLDER",
    "_V_ROW_WIDTH",
    "_V_SHEET_NAME",
    "_V_WIDTH",
    "_arb_rows",
    "_codes",
    "_expected_face",
    "_extra_keys",
    "_extract",
    "_header_violations",
    "_live",
    "_materialize",
    "_placeholder_positions",
    "_row_violations",
    "_spec",
    "_template_violations",
]


def test_anchor_worksheet_scope_and_width() -> None:
    """作业面规模、列数、单一列面 —— 三者塌陷时 ∀ 属性会恒真。"""
    codes = _codes()
    assert len(codes) == 16, (
        f"作业面实测 {len(codes)} 张（design 用户裁决 1：16 张，零 pending_manual）："
        f"{list(codes)}"
    )
    assert set(codes) == set(impl.X3_SHEET_SPECS), (
        "清单派生作业面与实现 `X3_SHEET_SPECS` 分叉：\n"
        f"  清单={sorted(codes)}\n  实现={sorted(impl.X3_SHEET_SPECS)}"
    )

    widths = {live.width for live in _live().templates.values()}
    assert len(widths) == 1, f"16 张登记的列跨度不一致：{sorted(widths)}"
    width = next(iter(widths))
    face = _expected_face()
    assert len(face) == width and width > 1, (
        f"实读列面 {len(face)} 项 != 登记跨度 {width}（或跨度退化）：{face}"
    )
    assert len(set(face)) == len(face), f"实读列面有重复标签：{face}"

    # 锚点（非属性）：实现的派生常量必须落在同一个期望值上，否则下面拿的是错期望值。
    # 「COLUMN_ORDER 是否真由清单派生」由 GS1 TestColumnOrderDerivedFromLedger 承载。
    assert tuple(impl.COLUMN_ORDER) == face, (
        "COLUMN_ORDER 与源模板实读列面不一致 ⇒ 本文件后续属性的期望值侧不可信：\n"
        f"  实读={face}\n  COLUMN_ORDER={tuple(impl.COLUMN_ORDER)}"
    )


def test_anchor_sheet_names_and_placeholder_distribution() -> None:
    """sheet 名非空 + 占位位分布 —— 证明「不假定只有 F 列」这一维真有实例。"""
    empty = [c for c in _codes() if not str(_spec(c).sheet_name).strip()]
    assert not empty, f"以下 sheet 的 catalog sheet_name 为空：{empty}（R3.8 期望值无从取得）"
    assert len({_spec(c).sheet_name for c in _codes()}) == len(_codes()), (
        "16 张的 sheet_name 不唯一 ⇒ 「产物 sheet 名 == 本张的 sheet_name」判据会互相蒙对"
    )

    dist: dict[int, list[str]] = {}
    for code in _codes():
        dist.setdefault(len(_placeholder_positions(code)), []).append(code)
    counts = {n: len(v) for n, v in sorted(dist.items())}
    assert counts == {1: 14, 2: 1, 6: 1}, (
        "占位位分布与实测不符（14 张 1 个 / L2-3 2 个 / N5-3 6 个）："
        f"{ {n: sorted(v) for n, v in sorted(dist.items())} }\n"
        "→ 该分布是「占位位按 field_keys 的 None 位判定、不假定只有 F 列」这一维的覆盖面证据"
    )
    assert _extra_keys(), "「多余字段」维度取不到任何键名 ⇒ 该维度空转"


def test_anchor_templates_dir_unchanged_after_reads() -> None:
    """R11.1：本文件打开的 16 个源模板，读取前后 ``(size, mtime_ns)`` 逐条未变。"""
    bundle = _live()
    assert len(bundle.before) == 16, f"只快照到 {len(bundle.before)} 个源模板，自检不成立"
    changed = {
        rel: (bundle.before[rel], bundle.after[rel])
        for rel in bundle.before
        if bundle.before[rel] != bundle.after[rel]
    }
    assert not changed, (
        "读取源模板后 backend/wp_templates/ 下这些文件发生变化（违反 R11.1；"
        "也可能是并发会话 / WPS 正在写该目录）：\n"
        + "".join(f"  {rel}: {b} → {a}\n" for rel, (b, a) in sorted(changed.items()))
    )


# ═══════════════════════════════════════════════════════════════════════════
# 5. P3-1 —— ∀ sheet × ∀ 行数据：产物列面与 sheet 名不变形（in-memory）
# ═══════════════════════════════════════════════════════════════════════════


def test_p3_1_export_column_face_invariant_over_any_rows() -> None:
    """∀ 16 张 sheet × ∀ 行数据：列头逐字 == 实读列面、sheet 名 == catalog、占位位留空保列序。

    覆盖的行数据维度：行数（含 0）· 字段缺失 · 多余字段 · ``None`` · 空串 · 长中文 ·
    特殊字符 · 极大/极小有限小数。**列面期望值与输入无关**，故任何「按行内容改列面」
    （如缺字段就少写一列）都会打红。
    """
    face = _expected_face()
    seen = {
        "sheets": set(),
        "row_counts": set(),
        "value_kinds": set(),
        "dropped": set(),
        "extras": set(),
        "placeholders": set(),
    }

    @given(rows=_arb_rows(len(face)))
    @_SETTINGS
    def prop(rows: list[_AbstractRow]) -> None:
        seen["row_counts"].add(len(rows))
        for row in rows:
            seen["dropped"].add(len(row.dropped))
            seen["extras"].add(len(row.extras))
            for value in row.values:
                if value is None:
                    seen["value_kinds"].add("none")
                elif isinstance(value, (int, float)):
                    seen["value_kinds"].add("number")
                elif value == "":
                    seen["value_kinds"].add("empty")
                elif len(value) > 30:
                    seen["value_kinds"].add("long_cn")
                elif any(ch in value for ch in "，。（）%&#@!?+-*/_|^$<>"):
                    seen["value_kinds"].add("special")
                else:
                    seen["value_kinds"].add("text")

        for code in _codes():
            spec = _spec(code)
            seen["sheets"].add(code)
            seen["placeholders"].add(len(_placeholder_positions(code)))
            payload = _materialize(code, rows)

            data_wb = impl.build_data_workbook(payload, code)
            data = _extract(data_wb, spec.sheet_name)
            problems = _header_violations(data, face, spec.sheet_name) + _row_violations(
                data, code, face, payload, loose=False
            )
            assert not problems, (
                f"{code}（{spec.sheet_name}）导出数据工作簿列面变形（{len(payload)} 行输入）：\n"
                + "".join(f"  {p}\n" for p in problems)
            )

            tpl = _extract(impl.build_template_workbook(code), spec.sheet_name)
            problems = (
                _header_violations(tpl, face, spec.sheet_name)
                + _row_violations(tpl, code, face, [], loose=False)
                + _template_violations(tpl, code)
            )
            assert not problems, (
                f"{code}（{spec.sheet_name}）导出空白模板列面变形：\n"
                + "".join(f"  {p}\n" for p in problems)
            )

    prop()

    assert seen["sheets"] == set(_codes()), (
        f"只覆盖了 {len(seen['sheets'])} 张 sheet：{sorted(seen['sheets'])}"
    )
    assert 0 in seen["row_counts"] and max(seen["row_counts"]) >= 3, (
        f"行数维度覆盖不足（须含 0 行与 ≥3 行）：{sorted(seen['row_counts'])}"
    )
    assert {"none", "empty", "number", "long_cn", "special", "text"} <= seen["value_kinds"], (
        f"取值维度覆盖不足：{sorted(seen['value_kinds'])}"
    )
    assert max(seen["dropped"]) >= 1, "「字段缺失」维度未覆盖"
    assert max(seen["extras"]) >= 1, "「多余字段」维度未覆盖"
    assert seen["placeholders"] == {1, 2, 6}, (
        f"占位位规模覆盖不足：{sorted(seen['placeholders'])}（须含 1 / 2 / 6）"
    )


# ═══════════════════════════════════════════════════════════════════════════
# 6. P3-2 —— 真实 xlsx 字节往返后列面与 sheet 名仍不变形
# ═══════════════════════════════════════════════════════════════════════════


def test_p3_2_column_face_survives_real_xlsx_bytes() -> None:
    """∀ 行数据 × 16 张：``wb.save()`` 出的**真实字节**再 load 回来，列面与 sheet 名不变。

    P3-1 判的是 workbook 对象，这条判的是用户真正拿到手的文件 —— 序列化会把空串写成
    空单元格、把 ``round(v, 2)`` 后的整值写成 int，判据据此放宽为 ``loose``（空 = None 或
    纯空白；文本按 strip 后逐字比；数值按相对误差 1e-6）。**列头、列序、列数、占位位
    这四项不放宽。**
    """
    face = _expected_face()
    seen_rows: set[int] = set()

    @given(rows=_arb_rows(len(face), max_rows=3))
    @settings(
        max_examples=_RUNS_ROUNDTRIP,
        deadline=None,
        suppress_health_check=[HealthCheck.too_slow],
    )
    def prop(rows: list[_AbstractRow]) -> None:
        seen_rows.add(len(rows))
        for code in _codes():
            spec = _spec(code)
            payload = _materialize(code, rows)
            buffer = io.BytesIO()
            impl.build_data_workbook(payload, code).save(buffer)
            raw = buffer.getvalue()
            assert raw[:2] == b"PK", f"{code}: 导出字节不是 zip 容器 ⇒ 不是有效 xlsx"
            reloaded = load_workbook(io.BytesIO(raw))
            try:
                product = _extract(reloaded, spec.sheet_name)
            finally:
                reloaded.close()
            problems = _header_violations(product, face, spec.sheet_name) + _row_violations(
                product, code, face, payload, loose=True
            )
            assert not problems, (
                f"{code}（{spec.sheet_name}）真实 xlsx 字节往返后列面变形"
                f"（{len(payload)} 行输入）：\n" + "".join(f"  {p}\n" for p in problems)
            )

    prop()
    assert 0 in seen_rows and max(seen_rows) >= 2, f"行数维度覆盖不足：{sorted(seen_rows)}"


# ═══════════════════════════════════════════════════════════════════════════
# 7. P3-3 —— 判据非退化：∀ 列面变形必被判错（反空转的正面证明）
# ═══════════════════════════════════════════════════════════════════════════

_MUTATIONS = (
    "header_drop",
    "header_extra",
    "header_swap",
    "header_rename",
    "header_dup",
    "sheet_rename",
    "row_skip_placeholder",
    "row_shift",
    "expected_perturb",
)

#: 每种变形**至少**要命中的违规类之一
_EXPECTED_TAGS: dict[str, tuple[str, ...]] = {
    "header_drop": (_V_MISSING, _V_WIDTH),
    "header_extra": (_V_EXTRA, _V_WIDTH),
    "header_swap": (_V_ORDER,),
    "header_rename": (_V_EXTRA, _V_MISSING),
    "header_dup": (_V_DUP, _V_MISSING),
    "sheet_rename": (_V_SHEET_NAME,),
    "row_skip_placeholder": (_V_ROW_WIDTH,),
    "row_shift": (_V_MISPLACED, _V_PLACEHOLDER),
    "expected_perturb": (_V_EXTRA, _V_MISSING, _V_ORDER, _V_WIDTH),
}


def _mutate(
    product: _Product,
    face: tuple[str, ...],
    code: str,
    kind: str,
    i: int,
    j: int,
) -> tuple[_Product, tuple[str, ...]]:
    """把真实产物（或期望值）改坏一处。返回 (变形后产物, 变形后期望列面)。"""
    title = product.title
    grid = [list(row) for row in product.grid]
    header = grid[0]
    expected = face
    placeholders = _placeholder_positions(code)

    if kind == "header_drop":
        del header[i]
    elif kind == "header_extra":
        header.insert(i, f"{_MARK}{i}")
    elif kind == "header_swap":
        header[i], header[j] = header[j], header[i]
    elif kind == "header_rename":
        header[i] = f"{header[i]}{_MARK}"
    elif kind == "header_dup":
        header[i] = header[j]
    elif kind == "sheet_rename":
        title = f"{title}{_MARK}"
    elif kind == "row_skip_placeholder":
        for pos in range(1, len(grid)):
            grid[pos] = [c for k, c in enumerate(grid[pos]) if k not in placeholders]
    elif kind == "row_shift":
        for pos in range(1, len(grid)):
            row = grid[pos]
            grid[pos] = row[-1:] + row[:-1]
    elif kind == "expected_perturb":
        expected = face[:i] + (f"{face[i]}{_MARK}",) + face[i + 1 :]
    else:  # pragma: no cover - 词表由 _MUTATIONS 限定
        raise AssertionError(kind)

    return _Product(title=title, grid=tuple(tuple(r) for r in grid)), expected


def _judge_one_mutation(
    *,
    product: _Product,
    payload: Sequence[Mapping[str, Any]],
    face: tuple[str, ...],
    code: str,
    kind: str,
    i: int,
    j: int,
) -> bool:
    """对同一份真实产物施加**一种**变形并判读；返回该实例是否**有效**（非恒等）。

    变形若退化成恒等（例如整行全空时 ``row_shift`` 前后相同、或该 sheet 无数据行时
    ``row_skip_placeholder`` 无作用），则反过来断言**仍是零违规** —— 不给判据留
    「反正总能报个错」的空子；非退化时断言「必被判错」且**违规类别对得上**。
    """
    if kind in ("header_swap", "header_dup") and i == j:
        j = (i + 1) % len(face)
    sheet_name = _spec(code).sheet_name
    broken, expected = _mutate(product, face, code, kind, i, j)
    problems = _header_violations(broken, expected, sheet_name) + _row_violations(
        broken, code, expected, payload, loose=False
    )
    if broken == product and expected == face:
        assert not problems, (
            f"{code} 的 {kind} 变形退化成恒等（产物与期望都没变），判据却报了违规 ⇒ "
            f"判据会无故打红：\n" + "".join(f"  {p}\n" for p in problems)
        )
        return False

    assert problems, (
        f"{code} 施加 {kind}(i={i}, j={j}) 变形后判据仍判通过 ⇒ **判据恒真（守卫缺陷）**。"
        f"\n  变形后列头={broken.grid[0] if broken.grid else ()}"
        f"\n  变形后表名={broken.title!r}\n  期望列面={expected}"
    )
    tags = _EXPECTED_TAGS[kind]
    assert any(tag in p for p in problems for tag in tags), (
        f"{code} 的 {kind} 变形被判错了，但违规类别不对（期望命中 {list(tags)}）：\n"
        + "".join(f"  {p}\n" for p in problems)
    )
    return True


def _sweep_all_mutations(
    *,
    code: str,
    rows: Sequence[_AbstractRow],
    face: tuple[str, ...],
    i: int,
    j: int,
    seen: dict[str, set[str]],
) -> None:
    """一次 CONTROL + **内循环遍历整张 ``_MUTATIONS`` 词表**（覆盖面不经随机抽样）。

    先跑 CONTROL（未变形产物 ⇒ 零违规），使「变形后打红」可归因于变形本身。产物只
    构造一次（``build_data_workbook`` + ``_extract`` 是这里唯一的重活），9 种变形共用
    它 ⇒ 把 ``kind`` 从 ``sampled_from`` 挪进内循环几乎不加成本。
    """
    spec = _spec(code)
    payload = _materialize(code, rows)
    product = _extract(impl.build_data_workbook(payload, code), spec.sheet_name)

    control = _header_violations(product, face, spec.sheet_name) + _row_violations(
        product, code, face, payload, loose=False
    )
    assert not control, (
        f"CONTROL 失败：{code} 未变形产物就被判违规 ⇒ 变形后的红不可归因：\n"
        + "".join(f"  {p}\n" for p in control)
    )

    seen["codes"].add(code)
    for kind in _MUTATIONS:
        seen["kinds"].add(kind)
        if _judge_one_mutation(
            product=product, payload=payload, face=face, code=code, kind=kind, i=i, j=j
        ):
            seen["effective"].add(kind)


def _nondegenerate_rows(width: int, count: int = 2) -> list[_AbstractRow]:
    """确定性扫描用的固定行：逐列取值互不相同且非空、无缺失字段。

    ``row_shift`` / ``row_skip_placeholder`` 在**空表或整行全同值**时会退化成恒等，
    退化实例进不了 ``seen["effective"]`` ⇒ 覆盖面若只靠 hypothesis 抽的 ``rows``，
    「某变形从未产生有效实例」这条断言仍是概率性的。用一组写死的非恒等行把它钉成确定的
    （占位位由 ``field_keys`` 的 ``None`` 位决定，天然无字段，故产物里仍是空串）。
    """
    return [
        _AbstractRow(
            values=tuple(f"{_MARK}{r}-{c}" for c in range(width)),
            dropped=frozenset(),
            extras=(),
        )
        for r in range(count)
    ]


def test_p3_3_judgement_rejects_any_deformed_column_face() -> None:
    """∀ 变形 × ∀ sheet × ∀ 行数据：判据必须报出对应类别的违规（证明判据不是恒真）。

    🔴 **变形维度是内循环遍历、不是 hypothesis 抽样**（任务 17.2 的修法）：``kind`` 曾是
    ``st.sampled_from(_MUTATIONS)``，覆盖断言 ``== set(_MUTATIONS)`` 因而把「hypothesis
    有没有把 9 个词各抽到过」当硬条件。本仓 hypothesis **6.152.4** 未 ``derandomize``、
    ``sampled_from`` 受 swarm testing 影响（单例会**整词禁用**一部分取值），150 次迭代
    漏抽某一词**真的会发生**，不是 uniform 的 ~2e-8：任务 6.2 首跑漏 ``sheet_rename``、
    Checkpoint 7 run1 漏 ``row_shift``（run2 自己转绿）；任务 17.1 的 13 文件套件内更是
    **连续两轮红在同一条同一词**、而单跑本文件全绿 ⇒ 还叠了同进程内其它 PBT 的全局状态
    （例库 / swarm 决策 / 固定用例顺序），不是纯随机。改成内循环后覆盖面**与抽样无关**，
    与 ``test_x3_parse_property`` / ``test_x3_keyfamily_property``「内循环遍历 16 张」同
    一范式。

    **不用 ``derandomize=True``**：那只是把随机换成固定种子，覆盖面依旧是抽样的副产物
    —— 换个 hypothesis 版本 / 换台机器 / 换个 swarm 决策照样漏，而且会把这条属性对
    ``rows`` 的探索能力一并锁死。也**不删**这两条覆盖断言、**不改**成 ``>=`` 阈值：
    它们正是防「变形维度空转」的那条。
    """
    face = _expected_face()
    width = len(face)
    seen: dict[str, set[str]] = {"kinds": set(), "effective": set(), "codes": set()}

    # ── 确定性扫描：∀16 张 × ∀9 变形 × 一组写死的非恒等行 ──────────────────────
    # 下面三条覆盖断言的实例**全部**由这一段供给 ⇒ 与 hypothesis 抽到什么无关。
    for code in _codes():
        _sweep_all_mutations(
            code=code, rows=_nondegenerate_rows(width), face=face, i=0, j=1, seen=seen
        )

    @given(
        rows=_arb_rows(width, max_rows=3),
        i=st.integers(min_value=0, max_value=width - 1),
        j=st.integers(min_value=0, max_value=width - 1),
        code=st.sampled_from(sorted(_codes())),
    )
    @_SETTINGS
    def prop(rows: list[_AbstractRow], i: int, j: int, code: str) -> None:
        # 随机侧照旧跑满 ∀ 行数据 × ∀ 变形位 × ∀ sheet（含 0 行 / 缺字段 / 多余字段 /
        # None / 长中文 / 极值小数），但它只**增补**覆盖，不再承担覆盖判据。
        _sweep_all_mutations(code=code, rows=rows, face=face, i=i, j=j, seen=seen)

    prop()
    assert seen["kinds"] == set(_MUTATIONS), (
        "变形词表未被内循环走遍（有分支被 continue/break 跳过、或词表被就地改过？）：缺 "
        f"{sorted(set(_MUTATIONS) - seen['kinds'])}"
    )
    assert seen["effective"] == set(_MUTATIONS), (
        "以下变形从未产生过有效（非恒等）实例 ⇒ 该维度实际空转（变形被写成恒等了？）："
        f"{sorted(set(_MUTATIONS) - seen['effective'])}"
    )
    assert seen["codes"] == set(_codes()), (
        f"16 张 sheet 未被走遍 ⇒ 确定性扫描被削过：缺 {sorted(set(_codes()) - seen['codes'])}"
    )


# ═══════════════════════════════════════════════════════════════════════════
# 8. P3-4 —— sheet 名真源 = ACNR catalog（替身驱动；值判据对该维恒绿）
# ═══════════════════════════════════════════════════════════════════════════


class _CatalogStandIn:
    """``CatalogIndex`` 的最小替身 —— 只提供实现真正用到的 ``sheets_by_code``。"""

    def __init__(self, sheets_by_code: dict[str, list[dict[str, Any]]]) -> None:
        self.sheets_by_code = sheets_by_code


@lru_cache(maxsize=1)
def _real_catalog_entries() -> dict[str, dict[str, Any]]:
    from app.services.acnr.catalog import get_catalog

    catalog = get_catalog()
    out: dict[str, dict[str, Any]] = {}
    for code in _codes():
        entries = catalog.sheets_by_code.get(code, [])
        assert len(entries) == 1, (
            f"{code}: catalog 里该 sheet_code 有 {len(entries)} 条条目 ⇒ sheet_name 不唯一"
        )
        out[code] = dict(entries[0])
    return out


def _standin(suffix: str) -> _CatalogStandIn:
    by_code: dict[str, list[dict[str, Any]]] = {}
    for code, entry in _real_catalog_entries().items():
        clone = dict(entry)
        clone["sheet_name"] = f"{entry['sheet_name']}{suffix}"
        by_code[code] = [clone]
    return _CatalogStandIn(by_code)


def _exec_clone_with_catalog(standin: _CatalogStandIn) -> tuple[types.ModuleType, int]:
    """用 catalog 替身重新 exec 一份实现副本，返回 (模块, ``get_catalog`` 被调次数)。

    🔴 ``exec_module`` 前先把克隆名注册进 ``sys.modules``、``finally`` 删 —— 与 GS1 的
    ``_exec_module_with_ledger`` 同一条 harness 纪律（GD-3：CPython 3.12 的
    ``dataclasses._process_class`` 会按 ``cls.__module__`` 反查 ``sys.modules``）。
    """
    clone_name = f"{_IMPL_NAME.rsplit('.', 1)[0]}._x3_column_property_clone"
    spec = importlib.util.spec_from_file_location(clone_name, _IMPL_PATH)
    assert spec is not None and spec.loader is not None, f"无法为 {_IMPL_PATH} 构造 import spec"
    module = importlib.util.module_from_spec(spec)
    sys.modules[clone_name] = module
    try:
        with mock.patch(
            "app.services.acnr.catalog.get_catalog", return_value=standin
        ) as patched:
            spec.loader.exec_module(module)
        return module, patched.call_count
    finally:
        sys.modules.pop(clone_name, None)


_arb_suffix = st.text(alphabet="替身探针ZQ0129", min_size=1, max_size=6)


def test_p3_4_sheet_name_really_comes_from_catalog() -> None:
    """∀ catalog 替身名：``X3_SHEET_SPECS[*].sheet_name`` 与**导出产物**的工作表名都必须跟着变。

    这是本文件相对 GS2 的实质增量。实测清单 ``provenance.column_source.tab`` 与 catalog
    ``sheet_name`` 对 16 张逐字相同 ⇒ 「导出名改从清单 tab 取」是值等价改动，一切值比对
    判据（含 GS2 那条）恒绿；只有替身驱动能抓住它。

    同时断言**列面不随 sheet 名替身而变**（两处真源互不串味：列面来自清单 ``column_map``，
    sheet 名来自 catalog）。
    """
    face = _expected_face()
    seen_suffixes: set[str] = set()

    @given(suffix=_arb_suffix)
    @settings(
        max_examples=_RUNS_STANDIN,
        deadline=None,
        suppress_health_check=[HealthCheck.too_slow],
    )
    def prop(suffix: str) -> None:
        seen_suffixes.add(suffix)
        module, hits = _exec_clone_with_catalog(_standin(suffix))
        assert hits >= len(_codes()), (
            f"模块装载期间只调了 {hits} 次 `get_catalog()`（期望 ≥ {len(_codes())}）⇒ "
            "sheet 名不可能逐张来自 catalog（写死一份、或从别处取）"
        )
        specs = getattr(module, "X3_SHEET_SPECS", None)
        assert isinstance(specs, dict) and len(specs) == len(_codes()), (
            f"替身装载后 X3_SHEET_SPECS 形态异常：{type(specs).__name__} / "
            f"{len(specs) if isinstance(specs, dict) else 'n/a'} 条"
        )

        real = _real_catalog_entries()
        offenders: list[str] = []
        for code in _codes():
            want = f"{real[code]['sheet_name']}{suffix}"
            got = getattr(specs[code], "sheet_name", None)
            if got != want:
                offenders.append(f"{code}: spec.sheet_name={got!r} 期望 {want!r}")
                continue
            for builder, label in (
                (module.build_template_workbook, "export-template"),
                (lambda c: module.build_data_workbook([], c), "export-data"),
            ):
                titles = list(builder(code).sheetnames)
                if not titles or titles[0] != want:
                    offenders.append(f"{code}: {label} 产物工作表名={titles} 期望首项 {want!r}")
        assert not offenders, (
            f"catalog 替身把 sheet_name 加了后缀 {suffix!r}，但以下未跟着变 ⇒ "
            "导出 sheet 名不是从 ACNR catalog 派生的（R3.8）：\n"
            + "".join(f"  {o}\n" for o in offenders)
        )

        assert tuple(getattr(module, "COLUMN_ORDER", ())) == face, (
            "换 catalog sheet 名把列面也带偏了 ⇒ 两处真源串味"
            f"（列面真源 = 清单 column_map）：{tuple(getattr(module, 'COLUMN_ORDER', ()))}"
        )

    prop()
    assert len(seen_suffixes) >= 5, f"替身名维度覆盖不足：{sorted(seen_suffixes)}"


def test_p3_4_reverse_selfcheck_real_names_have_no_probe_marker() -> None:
    """反向自检：真实 catalog 下 sheet 名与产物**不带**探针后缀（证明上一条不是恒真）。"""
    polluted = [c for c in _codes() if _MARK in _spec(c).sheet_name]
    assert not polluted, f"真实 sheet_name 里出现探针后缀 {polluted} ⇒ 替身未还原"
    for code in _codes():
        name = _spec(code).sheet_name
        titles = list(impl.build_template_workbook(code).sheetnames)
        assert titles and titles[0] == name, (
            f"{code}: 真实产物工作表名={titles} 期望首项 {name!r}"
        )
        assert _MARK not in titles[0]


# ═══════════════════════════════════════════════════════════════════════════
# 9. P3-5 —— 未登记 sheet 一律拒绝导出（列面判据的作业面边界）
# ═══════════════════════════════════════════════════════════════════════════


def test_p3_5_unregistered_sheet_never_yields_a_workbook() -> None:
    """∀ 不在作业面内的 sheet 码：两个导出出口都必须抛错，不得回退成「某张表的列面」。

    这条守的是 R3.1 / R3.8 的**边界**：若未登记 sheet 也能拿到 workbook，它的列头与
    sheet 名就来自某个兜底值 —— 那正是「源模板不存在的列」进入产物的路径（R3.6）。
    """
    registered = set(_codes())
    seen = 0

    @given(
        candidate=st.text(
            alphabet=string.ascii_uppercase + string.digits + "-", min_size=1, max_size=8
        )
    )
    @_SETTINGS
    def prop(candidate: str) -> None:
        nonlocal seen
        if candidate in registered:
            return
        seen += 1
        for builder, label in (
            (impl.build_template_workbook, "build_template_workbook"),
            (lambda c: impl.build_data_workbook([], c), "build_data_workbook"),
        ):
            with pytest.raises(impl.X3ContractError) as caught:
                builder(candidate)
            assert candidate in str(caught.value) or "作业面" in str(caught.value), (
                f"{label}({candidate!r}) 抛的错不可读：{caught.value}"
            )

    prop()
    assert seen >= 50, f"未登记 sheet 的有效实例只有 {seen} 个 ⇒ 该属性覆盖不足"

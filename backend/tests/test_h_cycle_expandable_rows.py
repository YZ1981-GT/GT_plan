"""H 类披露 sheet 可扩标记守卫（spec `h-cycle-…` Task 15 / R8.1~R8.5）

**裁决者 = `backend/wp_templates/H/*.xlsx`（openpyxl 直读）**，对照附注模板 JSON 行集。

## 三条不变式

1. **列头级 `……` 必须已展开** —— 作**列头**的省略号永远收不到数据
   （载荷按 `key: c.label` 推的是底稿实际分类名）⇒ JSON 的 `headers` / `columns[].label`
   里不得残留。实证 4 处列头标记（H1-listed `E13` · H5-listed `E7`/`F7` ·
   H7-listed 8 处 · H8-listed `E6`）都已展开。

2. **行级 `……` 按循环声明处置** —— 作**行**的省略号是源模板的**真实可扩位**，
   两种正确处置并存，故必须逐循环登记而不能一刀切：
   - `keep`：seed 进 JSON 行集且**行数与 xlsx 逐一相等**（H1/H7-listed/H8/H9）
   - `dynamic-rows`：改由「+ 新增类别」动态行承载、**有意不 seed**（H7-soe）
     —— 预置空 `……` 行会被同步推成「占位披露行」污染附注

3. **纯占位说明词不得 seed** —— `可无限量添加行` / `预留` / `可改名` 是**说明文字**
   不是数据行（H3-listed 有 7 处、H1-listed 有 2 处，JSON 里必须 0 处）。

🔴 `……` 的两种语义**处理方式相反**，这是本循环最易踩错的判断：
同一个符号作列头要**展开**、作行要**保留**（或显式改动态行）。
"""

from __future__ import annotations

import json
import warnings
from pathlib import Path

import pytest

warnings.filterwarnings("ignore", category=UserWarning, module="openpyxl")

import openpyxl

_ROOT = Path(__file__).resolve().parents[1]
_TPL_DIR = _ROOT / "wp_templates" / "H"
_DATA = _ROOT / "data"

ELLIPSIS = ("……", "…", "......")
PLACEHOLDER_WORDS = ("可无限量添加行", "预留", "可改名")


# ─────────────────── 逐循环声明（判据真源，每条带实证） ───────────────────

# (wp_code, variant) -> 附注章节号；None = 该变体无附注落点
SECTION_OF: dict[tuple[str, str], str | None] = {
    ("H1", "listed"): "五、22",
    ("H1", "soe"): "八、22",
    ("H2", "listed"): "五、23",
    ("H2", "soe"): "八、23",
    ("H3", "listed"): "五、21",
    ("H3", "soe"): "八、21",
    # H5 上市无附注落点（variant_matrix listed_* 均为 null，见 Task 11）
    ("H5", "listed"): None,
    ("H5", "soe"): "八、25",
    ("H7", "listed"): "五、24",
    ("H7", "soe"): "八、24",
    ("H8", "listed"): "五、25",
    ("H8", "soe"): "八、26",
    ("H9", "listed"): "五、47",
    ("H9", "soe"): "八、52",
}

# 行级 `……` 的处置策略：'keep'（seed 且行数相等）| 'dynamic-rows'（有意不 seed）
ELLIPSIS_ROW_POLICY: dict[tuple[str, str], tuple[str, str]] = {
    ("H1", "listed"): ("keep", "6 处（各层本期增加/减少下的可扩明细行）参与小计求和"),
    ("H7", "listed"): ("keep", "5 处（列转置表各层可扩明细行）参与小计"),
    ("H8", "listed"): ("keep", "6 处（三层各自本期增加/减少下的可扩行）参与 sumOf"),
    ("H9", "soe"): (
        "keep",
        "1 处（`附注披露信息（国企）!A11` 的可扩扣减行）—— 租赁负债国企侧五行表里的"
        "扣减明细可扩位，已 seed 进 八、52 行集",
    ),
    ("H7", "soe"): (
        "dynamic-rows",
        "国企侧 8 处 `……` 是「可继续加行」的行形态标记；H7 soe 已改动态类别行模型"
        "（ElMessageBox.prompt 输入类别名后建行），预置空 `……` 行会被同步推成"
        "「占位披露行」污染附注 ⇒ 有意不 seed（spec h7-biological-assets-disclosure-rebuild）",
    ),
}


def _visible_note_sheets(path: Path) -> list[tuple[str, openpyxl.worksheet.worksheet.Worksheet]]:
    wb = openpyxl.load_workbook(path, data_only=False)
    out = []
    for sn in wb.sheetnames:
        if "附注" not in sn:
            continue
        ws = wb[sn]
        if ws.sheet_state != "visible":
            continue
        out.append((sn, ws))
    return out


def _scan_marks(ws) -> dict[str, list[str]]:
    """区分**列头级**与**行级**省略号，以及纯占位说明词。

    判据：表头行（首列含「项」且含「目」的行）里的、或不在 A 列的省略号 = 列头级；
    其余（A 列数据区）= 行级。
    """
    header_rows: set[int] = set()
    for r in range(1, min(ws.max_row, 15) + 1):
        a = ws.cell(r, 1).value
        if a and "项" in str(a) and "目" in str(a):
            header_rows.add(r)

    res = {"header": [], "row": [], "placeholder": []}
    for row in ws.iter_rows():
        for c in row:
            if c.value is None:
                continue
            s = str(c.value).strip()
            if s in ELLIPSIS:
                key = "header" if (c.row in header_rows or c.column > 1) else "row"
                res[key].append(c.coordinate)
            elif len(s) <= 24 and any(w in s for w in PLACEHOLDER_WORDS):
                res["placeholder"].append(c.coordinate)
    return res


@pytest.fixture(scope="module")
def templates() -> dict[str, dict[str, dict[str, list[str]]]]:
    """`{wp_code: {variant: marks}}`。"""
    assert _TPL_DIR.exists(), f"H 模板目录不存在: {_TPL_DIR}"
    out: dict[str, dict[str, dict[str, list[str]]]] = {}
    for p in sorted(x for x in _TPL_DIR.glob("*.xlsx") if not x.name.startswith("~$")):
        code = p.name.split()[0]
        for sn, ws in _visible_note_sheets(p):
            variant = "listed" if "上市" in sn else "soe"
            out.setdefault(code, {})[variant] = _scan_marks(ws)
    return out


@pytest.fixture(scope="module")
def note_sections() -> dict[str, dict[str, dict]]:
    out: dict[str, dict[str, dict]] = {}
    for variant, fn in (("listed", "note_template_listed.json"), ("soe", "note_template_soe.json")):
        tpl = json.loads((_DATA / fn).read_text(encoding="utf-8"))
        secs = tpl.get("sections") if isinstance(tpl.get("sections"), list) else tpl
        out[variant] = {str(s.get("section_number") or ""): s for s in secs}
    return out


def _json_marks(sec: dict) -> dict[str, int]:
    """JSON 侧的省略号统计：行集里的 / headers 与 columns label 里的 / 占位词行。"""
    res = {"row": 0, "header": 0, "placeholder": 0}
    for t in sec.get("tables") or []:
        for rw in t.get("rows") or []:
            lab = str(rw.get("label") or "").strip()
            if lab in ELLIPSIS:
                res["row"] += 1
            elif any(w in lab for w in PLACEHOLDER_WORDS):
                res["placeholder"] += 1
        for h in t.get("headers") or []:
            if str(h).strip() in ELLIPSIS:
                res["header"] += 1
        for c in t.get("columns") or []:
            for k in ("label", "group"):
                if str(c.get(k) or "").strip() in ELLIPSIS:
                    res["header"] += 1
    return res


# ────────────────────────── 解析器自检 ──────────────────────────


def test_scanner_self_check(templates):
    """防「扫描面为空 / 分类器失效 → 全部断言空转」。"""
    assert len(templates) >= 9, f"H 模板扫到 {len(templates)} 个，疑似 glob 失效"
    # 已知锚点（本轮实证）：列头级与行级都必须真被区分出来
    assert templates["H8"]["listed"]["header"] == ["E6"], templates["H8"]["listed"]
    assert len(templates["H8"]["listed"]["row"]) == 6, templates["H8"]["listed"]
    assert len(templates["H3"]["listed"]["placeholder"]) == 7, templates["H3"]["listed"]
    # 反向：分类器不能把所有标记都塞进同一桶
    assert templates["H7"]["listed"]["header"] and templates["H7"]["listed"]["row"]


def test_policy_table_covers_all_row_marks(templates):
    """凡 xlsx 有行级 `……` 的 (循环, 变体) 都必须在策略表里登记（防漏判）。"""
    missing = []
    for code, variants in templates.items():
        for variant, marks in variants.items():
            if marks["row"] and (code, variant) not in ELLIPSIS_ROW_POLICY:
                missing.append(f"{code}-{variant}({len(marks['row'])} 处)")
    assert missing == [], (
        "以下 (循环, 变体) 的源 sheet 有行级可扩标记但未登记处置策略："
        + ", ".join(missing)
    )


def test_policy_table_has_no_stale_entries(templates):
    """策略表不得有 stale 条目（xlsx 已无该标记却还登记着）。"""
    stale = []
    for (code, variant), (kind, _reason) in ELLIPSIS_ROW_POLICY.items():
        marks = (templates.get(code) or {}).get(variant)
        if marks is None or not marks["row"]:
            stale.append(f"{code}-{variant}({kind})")
    assert stale == [], "策略表条目已过期（源 sheet 无行级标记）：" + ", ".join(stale)


def test_every_policy_entry_has_reason():
    for (code, variant), (kind, reason) in ELLIPSIS_ROW_POLICY.items():
        assert kind in ("keep", "dynamic-rows"), f"{code}-{variant} 策略取值非法: {kind}"
        assert len(reason.strip()) >= 15, f"{code}-{variant} 理由过短（防当逃逸阀用）"


# ────────────────── 不变式 1：列头级 `……` 必须已展开 ──────────────────


@pytest.mark.parametrize(("code", "variant"), sorted(SECTION_OF))
def test_header_ellipsis_expanded(templates, note_sections, code, variant):
    secnum = SECTION_OF[(code, variant)]
    if secnum is None:
        pytest.skip(f"{code}-{variant} 无附注落点")
    marks = (templates.get(code) or {}).get(variant)
    if marks is None:
        pytest.skip(f"{code}-{variant} 源 sheet 不存在")
    sec = note_sections[variant].get(secnum)
    assert sec is not None, f"{variant} 缺章节 {secnum!r}"
    js = _json_marks(sec)
    assert js["header"] == 0, (
        f"{code}-{variant} ({secnum}) 的 headers/columns 残留 {js['header']} 个省略号列头 —— "
        f"列头级 `……` 永远收不到数据，必须展开为实际分类名"
        f"（源 sheet 列头标记位置: {marks['header']}）"
    )


# ────────────────── 不变式 2：行级 `……` 按策略处置 ──────────────────


@pytest.mark.parametrize(("code", "variant"), sorted(ELLIPSIS_ROW_POLICY))
def test_row_ellipsis_matches_policy(templates, note_sections, code, variant):
    kind, reason = ELLIPSIS_ROW_POLICY[(code, variant)]
    secnum = SECTION_OF[(code, variant)]
    assert secnum is not None, f"{code}-{variant} 登记了策略却无附注落点"
    xlsx_n = len(templates[code][variant]["row"])
    sec = note_sections[variant].get(secnum)
    assert sec is not None, f"{variant} 缺章节 {secnum!r}"
    js_n = _json_marks(sec)["row"]

    if kind == "keep":
        assert js_n == xlsx_n, (
            f"{code}-{variant} ({secnum}) 行级可扩位数量不符：源 sheet {xlsx_n} 处 "
            f"vs 附注模板 {js_n} 处。这些 `……` 是**真实可扩行**且参与小计求和，"
            f"删掉会让审计师无处添加明细。理由登记: {reason}"
        )
    else:  # dynamic-rows
        assert js_n == 0, (
            f"{code}-{variant} ({secnum}) 已声明改动态行承载，附注模板不得 seed "
            f"`……` 空行（实为 {js_n} 处）—— 预置空占位会被推成「占位披露行」污染附注。"
            f"理由: {reason}"
        )


# ────────────────── 不变式 3：纯占位说明词不得 seed ──────────────────


@pytest.mark.parametrize(("code", "variant"), sorted(SECTION_OF))
def test_placeholder_words_not_seeded(templates, note_sections, code, variant):
    secnum = SECTION_OF[(code, variant)]
    if secnum is None:
        pytest.skip(f"{code}-{variant} 无附注落点")
    marks = (templates.get(code) or {}).get(variant)
    if marks is None:
        pytest.skip(f"{code}-{variant} 源 sheet 不存在")
    sec = note_sections[variant].get(secnum)
    assert sec is not None
    js = _json_marks(sec)
    assert js["placeholder"] == 0, (
        f"{code}-{variant} ({secnum}) seed 了 {js['placeholder']} 个占位说明行 —— "
        f"「可无限量添加行」/「预留」/「可改名」是**说明文字**不是数据行，"
        f"语义应移入 `guidance`（源 sheet 占位位置: {marks['placeholder']}）"
    )


# ────────────────── 反向自检（证明断言不是空转） ──────────────────


def test_selfcheck_header_ellipsis_would_be_detected():
    fake = {"tables": [{"headers": ["项目", "……", "合计"], "rows": [], "columns": []}]}
    assert _json_marks(fake)["header"] == 1


def test_selfcheck_column_label_ellipsis_would_be_detected():
    fake = {"tables": [{"headers": [], "rows": [], "columns": [{"label": "……"}]}]}
    assert _json_marks(fake)["header"] == 1


def test_selfcheck_row_ellipsis_counted_separately():
    fake = {
        "tables": [
            {
                "headers": ["项目"],
                "columns": [],
                "rows": [{"label": "……"}, {"label": "可无限量添加行"}, {"label": "正常行"}],
            }
        ]
    }
    got = _json_marks(fake)
    assert got == {"row": 1, "header": 0, "placeholder": 1}, got


def test_selfcheck_classifier_distinguishes_header_and_row(templates):
    """同一个符号在列头与数据区被分到不同桶（这是本守卫的核心区分）。"""
    h8 = templates["H8"]["listed"]
    assert h8["header"] and h8["row"]
    assert set(h8["header"]).isdisjoint(set(h8["row"]))

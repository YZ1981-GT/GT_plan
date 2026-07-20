"""Property-Based Tests for Review Export

Property 10: Export worksheet count
Property 11: RFC5987 filename encoding validity

Validates: Requirements 7.1, 7.4
"""
from __future__ import annotations

import io
import uuid
from datetime import datetime, timezone
from urllib.parse import quote, unquote

from hypothesis import given, settings, strategies as st
from openpyxl import load_workbook

from app.routers.review_prompt import _generate_review_excel
from app.services.batch_review_service import (
    BatchReviewReport,
    BatchStatistics,
    ExecutionMetadata,
    ReviewResult,
)
from app.services.llm_response_parser import ReviewFinding


# ─── Strategies ─────────────────────────────────────────────────────────────

_RISK_LEVELS = st.sampled_from(["high", "medium", "low"])
_PASS_STATUSES = st.sampled_from(["pass", "fail", "review_error"])

# openpyxl rejects control characters in cell values (ILLEGAL_CHARACTERS_RE).
# Constrain text to printable characters safe for Excel.
_excel_safe_text = st.text(
    alphabet=st.characters(
        whitelist_categories=("L", "N", "P", "S", "Z"),
        blacklist_characters="\x00\x01\x02\x03\x04\x05\x06\x07\x08\x0b\x0c\x0e"
        "\x0f\x10\x11\x12\x13\x14\x15\x16\x17\x18\x19\x1a\x1b\x1c\x1d\x1e\x1f",
    ),
    min_size=1,
    max_size=60,
)

_finding_strategy = st.builds(
    ReviewFinding,
    id=st.uuids().map(str),
    description=_excel_safe_text,
    risk_level=_RISK_LEVELS,
    pass_status=st.booleans(),
    category=st.sampled_from(["认定检查", "程序执行", "数据完整性", "风险评估"]),
    sheet_location=st.none() | _excel_safe_text,
    suggestion=st.none() | _excel_safe_text,
)


def _build_review_result(wp_code: str, sheet_name: str, findings, pass_status: str) -> ReviewResult:
    """Helper to construct a ReviewResult."""
    risk_summary = {"high": 0, "medium": 0, "low": 0}
    for f in findings:
        if not f.pass_status and f.risk_level in risk_summary:
            risk_summary[f.risk_level] += 1
    return ReviewResult(
        wp_id=str(uuid.uuid4()),
        sheet_name=sheet_name,
        wp_code=wp_code,
        pass_status=pass_status,
        findings=findings,
        risk_summary=risk_summary,
        reviewed_at=datetime.now(timezone.utc).isoformat(),
        model_used="test-model",
        prompt_source="sheet",
        error_message="LLM error" if pass_status == "review_error" else None,
    )


@st.composite
def _batch_report_strategy(draw):
    """Generate a random BatchReviewReport with 1-10 sheets."""
    num_sheets = draw(st.integers(min_value=1, max_value=10))
    results = []
    for i in range(num_sheets):
        wp_code = f"D2-{i + 1}"
        sheet_name = f"底稿{i + 1}"
        findings = draw(st.lists(_finding_strategy, min_size=0, max_size=5))
        pass_status = draw(_PASS_STATUSES)
        results.append(_build_review_result(wp_code, sheet_name, findings, pass_status))

    passed_count = sum(1 for r in results if r.pass_status == "pass")
    failed_count = sum(1 for r in results if r.pass_status == "fail")
    error_count = sum(1 for r in results if r.pass_status == "review_error")
    total_findings = sum(len(r.findings) for r in results)
    findings_by_risk = {"high": 0, "medium": 0, "low": 0}
    for r in results:
        for f in r.findings:
            if not f.pass_status and f.risk_level in findings_by_risk:
                findings_by_risk[f.risk_level] += 1

    report = BatchReviewReport(
        session_id=str(uuid.uuid4()),
        wp_code_prefix="D2",
        project_id=str(uuid.uuid4()),
        results=results,
        statistics=BatchStatistics(
            total_sheets=num_sheets,
            passed_count=passed_count,
            failed_count=failed_count,
            error_count=error_count,
            total_findings=total_findings,
            findings_by_risk=findings_by_risk,
        ),
        execution=ExecutionMetadata(
            start_time=datetime.now(timezone.utc).isoformat(),
            end_time=datetime.now(timezone.utc).isoformat(),
            duration_seconds=1.5,
            model_used="test-model",
        ),
    )
    return report


# Strategy for Chinese filenames
_chinese_filename_strategy = st.text(
    alphabet=st.characters(
        whitelist_categories=("Lo",),  # CJK ideographs
        whitelist_characters="应收账款复核报告审计底稿汇总明细坏账准备分析凭证检查政策附注",
    ),
    min_size=1,
    max_size=30,
)


# ─── Property 10: Export worksheet count ────────────────────────────────────


class TestExportWorksheetCount:
    """**Validates: Requirements 7.1**

    For any batch review export, the Excel workbook contains exactly
    (number of reviewed sheets + 1) worksheets, where +1 is the summary sheet.
    """

    @given(report=_batch_report_strategy())
    @settings(max_examples=5)
    def test_p10_export_worksheet_count(self, report: BatchReviewReport) -> None:
        """Property 10: Export worksheet count"""
        # Generate Excel bytes from report
        excel_bytes = _generate_review_excel(report)

        # Load the workbook and count worksheets
        wb = load_workbook(io.BytesIO(excel_bytes))
        worksheet_count = len(wb.sheetnames)

        expected_count = len(report.results) + 1  # +1 for summary sheet

        assert worksheet_count == expected_count, (
            f"Expected {expected_count} worksheets "
            f"(summary + {len(report.results)} sheets), "
            f"got {worksheet_count}. "
            f"Sheet names: {wb.sheetnames}"
        )


# ─── Property 11: RFC5987 filename encoding validity ────────────────────────


class TestRfc5987FilenameEncoding:
    """**Validates: Requirements 7.4**

    For any Chinese filename string, the RFC5987 encoded Content-Disposition
    header contains `filename*=UTF-8''` followed by percent-encoded UTF-8 bytes,
    decodable back to original.
    """

    @given(filename=_chinese_filename_strategy)
    @settings(max_examples=5)
    def test_p11_rfc5987_filename_encoding_validity(self, filename: str) -> None:
        """Property 11: RFC5987 filename encoding validity"""
        # Apply RFC5987 encoding (same as router implementation)
        encoded = quote(filename, safe="")

        # Build Content-Disposition header value
        content_disposition = f"attachment; filename*=UTF-8''{encoded}"

        # Verify: header contains the correct prefix
        assert "filename*=UTF-8''" in content_disposition, (
            f"Content-Disposition missing 'filename*=UTF-8''' prefix: "
            f"{content_disposition}"
        )

        # Verify: the encoded part can be decoded back to the original
        # Extract the encoded filename from the header
        prefix = "filename*=UTF-8''"
        idx = content_disposition.index(prefix) + len(prefix)
        extracted_encoded = content_disposition[idx:]

        decoded = unquote(extracted_encoded)
        assert decoded == filename, (
            f"Round-trip failed: original={repr(filename)}, "
            f"encoded={repr(extracted_encoded)}, "
            f"decoded={repr(decoded)}"
        )

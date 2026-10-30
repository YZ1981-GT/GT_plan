"""批量建项与合并树导入收口回归。"""
from __future__ import annotations

from io import BytesIO
from types import SimpleNamespace
from uuid import UUID, uuid4

import pytest
import pytest_asyncio
from openpyxl import Workbook, load_workbook
from sqlalchemy import func, select
from sqlalchemy.ext.asyncio import AsyncSession, async_sessionmaker, create_async_engine

import app.models  # noqa: F401 — 注册全部 ORM
from app.models.base import Base
from app.models.core import Project
from app.services.batch_project_execution import (
    apply_validation_plan,
    create_validation_plan,
)
from app.services.batch_project_service import (
    _TEMPLATE_COLUMNS,
    generate_template,
    parse_and_import,
    validate_batch,
)

ROOT = "91110000100000000R"
PARENT = "911100002000000005"
CHILD = "91110000300000000G"
BRANCH = "91110000400000000U"


@pytest_asyncio.fixture
async def engine():
    from sqlalchemy.dialects.sqlite.base import SQLiteTypeCompiler
    SQLiteTypeCompiler.visit_JSONB = SQLiteTypeCompiler.visit_JSON
    if not hasattr(SQLiteTypeCompiler, "visit_ARRAY"):
        SQLiteTypeCompiler.visit_ARRAY = lambda self, type_, **kw: "TEXT"
    eng = create_async_engine("sqlite+aiosqlite:///:memory:")
    async with eng.begin() as conn:
        await conn.run_sync(Base.metadata.create_all)
    yield eng
    await eng.dispose()


@pytest_asyncio.fixture
async def db(engine):
    factory = async_sessionmaker(engine, class_=AsyncSession, expire_on_commit=False)
    async with factory() as session:
        yield session


def make_excel(rows: list[list], *, columns: list[str] | None = None) -> bytes:
    workbook = Workbook()
    sheet = workbook.active
    sheet.title = "数据"
    sheet.append(columns or _TEMPLATE_COLUMNS)
    for row in rows:
        sheet.append(row)
    output = BytesIO()
    workbook.save(output)
    return output.getvalue()


def row(name: str, code: str, year: int | float, scope: str = "单户", *,
        parent: str = "", ultimate: str = "", relation: str = "") -> list:
    return [name, code, name[:20], year, "年报审计", "企业会计准则", scope,
            parent, ultimate, relation, "新建", "", ""]


def admin():
    return SimpleNamespace(id=uuid4(), role=SimpleNamespace(value="admin"))

@pytest.mark.asyncio
async def test_template_has_empty_data_sheet_and_separate_examples():
    output = await generate_template(mode="consolidation")
    workbook = load_workbook(output)
    assert not any(
        any(value not in (None, "") for value in row_values)
        for row_values in workbook["数据"].iter_rows(min_row=2, values_only=True)
    )
    assert workbook["填写示例"].max_row == 5
    assert [cell.value for cell in workbook["数据"][1]] == _TEMPLATE_COLUMNS
    validations = list(workbook["数据"].data_validations.dataValidation)
    assert len(validations) == 5
    assert workbook["_GT_META"]["B1"].value == "2"
    assert workbook["_GT_META"]["B2"].value == "consolidation"


@pytest.mark.asyncio
async def test_decimal_year_is_rejected_without_truncation(db):
    result = await validate_batch(make_excel([row("甲公司", CHILD, 2025.5)]), db)
    assert result.valid is False
    assert "不能包含小数" in "；".join(result.errors[0].errors)


@pytest.mark.asyncio
async def test_same_company_year_two_scopes_are_not_batch_duplicates(db):
    rows = [
        row("集团", ROOT, 2025, "单户", ultimate=ROOT),
        row("集团", ROOT, 2025, "合并", ultimate=ROOT),
    ]
    result = await validate_batch(make_excel(rows), db)
    assert result.valid is True, result.errors
    assert result.summary.create == 2


@pytest.mark.asyncio
async def test_database_existing_key_is_reported_during_validation(db):
    from app.models.audit_platform_schemas import BasicInfoSchema
    from app.services.project_wizard_service import create_project
    await create_project(BasicInfoSchema(
        client_name="甲公司", audit_year=2025, project_type="annual",
        accounting_standard="enterprise", company_code=CHILD,
        short_name="甲", report_scope="standalone",
    ), db)
    result = await validate_batch(make_excel([row("甲公司", CHILD, 2025)]), db)
    assert result.valid is False
    assert "当前不能新建" in "；".join(result.errors[0].errors)


@pytest.mark.asyncio
async def test_same_ultimate_across_years_creates_one_root_per_year(engine):
    factory = async_sessionmaker(engine, class_=AsyncSession, expire_on_commit=False)
    rows = [
        row("子公司2025", CHILD, 2025, parent=ROOT, ultimate=ROOT, relation="子公司"),
        row("子公司2026", CHILD, 2026, parent=ROOT, ultimate=ROOT, relation="子公司"),
    ]
    async with factory() as session:
        result = await parse_and_import(make_excel(rows), session)
        assert result.success_count == 2
        assert result.auto_root_count == 2
    async with factory() as session:
        count = await session.scalar(select(func.count()).select_from(Project).where(
            Project.company_code == ROOT,
            Project.report_scope == "consolidated",
            Project.is_deleted == False,  # noqa: E712
        ))
        assert count == 2

@pytest.mark.asyncio
async def test_persisted_plan_applies_once_and_replays_receipt(engine):
    factory = async_sessionmaker(engine, class_=AsyncSession, expire_on_commit=False)
    actor = admin()
    file_bytes = make_excel([row("独立企业", BRANCH, 2027)])
    key = "plan-idempotency-0001"
    async with factory() as session:
        plan, validation, reused = await create_validation_plan(
            file_bytes, session, actor=actor, file_name="批量.xlsx",
            mode="project", idempotency_key=key,
        )
        assert validation.valid is True
        assert reused is False
        plan_id = plan.id
        await session.commit()
    async with factory() as session:
        first = await apply_validation_plan(
            session, plan_id=plan_id, actor=actor, idempotency_key=key
        )
        await session.commit()
        assert first.created_count == 1
        assert first.idempotent is False
    async with factory() as session:
        replay = await apply_validation_plan(
            session, plan_id=plan_id, actor=actor, idempotency_key=key
        )
        await session.commit()
        assert replay.created_count == 1
        assert replay.idempotent is True
        count = await session.scalar(select(func.count()).select_from(Project).where(
            Project.company_code == BRANCH,
            Project.audit_year == 2027,
            Project.is_deleted == False,  # noqa: E712
        ))
        assert count == 1


@pytest.mark.asyncio
async def test_same_idempotency_key_different_content_is_conflict(db):
    actor = admin()
    await create_validation_plan(
        make_excel([row("甲公司", CHILD, 2028)]), db,
        actor=actor, file_name="a.xlsx", mode="project",
        idempotency_key="same-key-00000001",
    )
    await db.commit()
    with pytest.raises(Exception) as captured:
        await create_validation_plan(
            make_excel([row("乙公司", BRANCH, 2028)]), db,
            actor=actor, file_name="b.xlsx", mode="project",
            idempotency_key="same-key-00000001",
        )
    assert getattr(captured.value, "status_code", None) == 409


@pytest.mark.asyncio
async def test_legacy_partial_failure_does_not_rollback_prior_success(db, monkeypatch):
    """故障注入在下一层 flush：第二个 Project flush 失败，第一行仍真实落库且回执一致。"""
    original_flush = db.flush
    project_flushes = 0

    async def failing_flush(*args, **kwargs):
        nonlocal project_flushes
        pending_projects = [item for item in db.new if isinstance(item, Project)]
        if pending_projects:
            project_flushes += 1
            if project_flushes == 2:
                raise RuntimeError("injected project flush failure")
        return await original_flush(*args, **kwargs)

    monkeypatch.setattr(db, "flush", failing_flush)
    result = await parse_and_import(make_excel([
        row("甲公司", CHILD, 2029),
        row("乙公司", BRANCH, 2029),
    ]), db)
    assert result.success_count == 1
    assert result.fail_count == 1
    count = await db.scalar(select(func.count()).select_from(Project).where(
        Project.audit_year == 2029, Project.is_deleted == False,  # noqa: E712
    ))
    assert count == 1

@pytest.mark.asyncio
async def test_group_export_edit_validate_apply_updates_relationship(engine):
    from app.models.audit_platform_schemas import BasicInfoSchema
    from app.services.batch_project_export import export_group_tree
    from app.services.project_wizard_service import create_project
    factory = async_sessionmaker(engine, class_=AsyncSession, expire_on_commit=False)
    async with factory() as session:
        root_project = await create_project(BasicInfoSchema(
            client_name="集团", audit_year=2030, project_type="annual",
            accounting_standard="enterprise", company_code=ROOT,
            short_name="集团", report_scope="consolidated",
            ultimate_company_code=ROOT,
        ), session)
        await create_project(BasicInfoSchema(
            client_name="中间控股", audit_year=2030, project_type="annual",
            accounting_standard="enterprise", company_code=PARENT,
            short_name="中间", report_scope="standalone",
            parent_company_code=ROOT, ultimate_company_code=ROOT,
            relation_to_parent="subsidiary",
        ), session)
        await create_project(BasicInfoSchema(
            client_name="子公司", audit_year=2030, project_type="annual",
            accounting_standard="enterprise", company_code=CHILD,
            short_name="子公司", report_scope="standalone",
            parent_company_code=ROOT, ultimate_company_code=ROOT,
            relation_to_parent="subsidiary",
        ), session)
        await create_project(BasicInfoSchema(
            client_name="子公司", audit_year=2030, project_type="annual",
            accounting_standard="enterprise", company_code=CHILD,
            short_name="子公司合并", report_scope="consolidated",
            parent_company_code=ROOT, ultimate_company_code=ROOT,
            relation_to_parent="subsidiary",
        ), session)
        output, ids = await export_group_tree(root_project.id, session)
        assert root_project.id in ids

    workbook = load_workbook(output)
    sheet = workbook["数据"]
    assert "说明事项" in workbook.sheetnames
    assert sheet.column_dimensions["L"].hidden is True
    assert sheet.column_dimensions["M"].hidden is True
    assert sheet["L2"].protection.locked is True
    assert sheet["M2"].protection.locked is True
    child_row = next(cells for cells in sheet.iter_rows(min_row=2) if cells[1].value == CHILD)
    child_row[7].value = PARENT
    edited = BytesIO()
    workbook.save(edited)

    actor = admin()
    key = "group-update-plan-0001"
    async with factory() as session:
        plan, validation, _ = await create_validation_plan(
            edited.getvalue(), session, actor=actor, file_name="集团树.xlsx",
            mode="consolidation", idempotency_key=key,
        )
        assert validation.valid is True, validation.errors
        assert validation.summary.update == 1
        assert validation.summary.unchanged == 2
        plan_id = plan.id
        await session.commit()
    async with factory() as session:
        result = await apply_validation_plan(
            session, plan_id=plan_id, actor=actor, idempotency_key=key
        )
        await session.commit()
        assert result.updated_count == 1
        children = list((await session.execute(select(Project).where(
            Project.company_code == CHILD, Project.audit_year == 2030,
        ))).scalars().all())
        assert len(children) == 2
        assert {child.parent_company_code for child in children} == {PARENT}

@pytest.mark.asyncio
async def test_manager_creator_gets_assignment_and_edit_membership(db):
    from app.models.base import PermissionLevel, ProjectUserRole, UserRole
    from app.models.core import ProjectUser, User
    from app.models.staff_models import ProjectAssignment, StaffMember
    manager = User(
        username="batch_manager", email="batch_manager@example.test",
        hashed_password="x", role=UserRole.manager,
    )
    db.add(manager)
    await db.flush()
    staff = StaffMember(user_id=manager.id, name="批量项目经理", employee_no="BP-MGR-1")
    db.add(staff)
    await db.commit()

    result = await parse_and_import(
        make_excel([row("经理创建项目", CHILD, 2031)]), db,
        current_user=manager,
    )
    assert result.success_count == 1
    project = (await db.execute(select(Project).where(
        Project.company_code == CHILD, Project.audit_year == 2031,
    ))).scalar_one()
    assert project.manager_id == manager.id
    member = (await db.execute(select(ProjectUser).where(
        ProjectUser.project_id == project.id,
        ProjectUser.user_id == manager.id,
        ProjectUser.is_deleted == False,  # noqa: E712
    ))).scalar_one()
    assert member.role == ProjectUserRole.manager
    assert member.permission_level == PermissionLevel.edit
    assignment = (await db.execute(select(ProjectAssignment).where(
        ProjectAssignment.project_id == project.id,
        ProjectAssignment.staff_id == staff.id,
        ProjectAssignment.is_deleted == False,  # noqa: E712
    ))).scalar_one()
    assert assignment.role == "manager"

@pytest.mark.parametrize(
    ("role", "status"),
    [("admin", 200), ("partner", 200), ("manager", 200),
     ("auditor", 403), ("qc", 403), ("eqcr", 403), ("readonly", 403)],
)
def test_batch_template_role_gate(role: str, status: int):
    from fastapi import FastAPI
    from fastapi.testclient import TestClient
    from app.deps import get_current_user
    from app.routers.batch_project import router
    app = FastAPI()
    app.include_router(router)
    app.dependency_overrides[get_current_user] = lambda: SimpleNamespace(
        id=uuid4(), role=SimpleNamespace(value=role)
    )
    response = TestClient(app).get("/api/projects/batch-template")
    assert response.status_code == status


@pytest.mark.asyncio
async def test_batch_export_rejects_mixed_authorized_ids(db):
    from app.models.base import PermissionLevel, ProjectUserRole, UserRole
    from app.models.core import ProjectUser, User
    from app.routers.batch_project import BatchExportRequest, batch_export_projects
    user = User(
        username="exporter", email="exporter@example.test",
        hashed_password="x", role=UserRole.manager,
    )
    db.add(user)
    await db.flush()
    one = Project(name="可读_2032", client_name="可读", company_code=CHILD,
                  short_name="可读", audit_year=2032, report_scope="standalone")
    two = Project(name="不可读_2032", client_name="不可读", company_code=BRANCH,
                  short_name="不可读", audit_year=2032, report_scope="standalone")
    db.add_all([one, two])
    await db.flush()
    db.add(ProjectUser(
        project_id=one.id, user_id=user.id,
        role=ProjectUserRole.manager, permission_level=PermissionLevel.edit,
    ))
    await db.commit()
    with pytest.raises(Exception) as captured:
        await batch_export_projects(
            BatchExportRequest(project_ids=[one.id, two.id]), db, user
        )
    assert getattr(captured.value, "status_code", None) == 403

@pytest.mark.asyncio
async def test_batch_reparent_validates_final_graph_not_intermediate_order(engine):
    from app.models.audit_platform_schemas import BasicInfoSchema
    from app.services.batch_project_export import export_group_tree
    from app.services.project_wizard_service import create_project
    factory = async_sessionmaker(engine, class_=AsyncSession, expire_on_commit=False)
    async with factory() as session:
        root_project = await create_project(BasicInfoSchema(
            client_name="换序集团", audit_year=2033, project_type="annual",
            accounting_standard="enterprise", company_code=ROOT,
            short_name="换序集团", report_scope="consolidated",
            ultimate_company_code=ROOT,
        ), session)
        await create_project(BasicInfoSchema(
            client_name="甲公司", audit_year=2033, project_type="annual",
            accounting_standard="enterprise", company_code=PARENT,
            short_name="甲", report_scope="standalone",
            parent_company_code=ROOT, ultimate_company_code=ROOT,
            relation_to_parent="subsidiary",
        ), session)
        await create_project(BasicInfoSchema(
            client_name="乙公司", audit_year=2033, project_type="annual",
            accounting_standard="enterprise", company_code=CHILD,
            short_name="乙", report_scope="standalone",
            parent_company_code=PARENT, ultimate_company_code=ROOT,
            relation_to_parent="subsidiary",
        ), session)
        output, _ = await export_group_tree(root_project.id, session)

    workbook = load_workbook(output)
    sheet = workbook["数据"]
    rows_by_code = {cells[1].value: cells for cells in sheet.iter_rows(min_row=2)}
    rows_by_code[PARENT][7].value = CHILD
    rows_by_code[CHILD][7].value = ROOT
    edited = BytesIO()
    workbook.save(edited)
    actor = admin()
    key = "final-graph-order-0001"
    async with factory() as session:
        plan, validation, _ = await create_validation_plan(
            edited.getvalue(), session, actor=actor, file_name="换序.xlsx",
            mode="consolidation", idempotency_key=key,
        )
        assert validation.valid is True, validation.errors
        assert validation.summary.update == 2
        plan_id = plan.id
        await session.commit()
    async with factory() as session:
        result = await apply_validation_plan(
            session, plan_id=plan_id, actor=actor, idempotency_key=key
        )
        await session.commit()
        assert result.updated_count == 2
        projects = {
            project.company_code: project
            for project in (await session.execute(select(Project).where(
                Project.audit_year == 2033,
                Project.company_code.in_([PARENT, CHILD]),
            ))).scalars().all()
        }
        assert projects[PARENT].parent_company_code == CHILD
        assert projects[CHILD].parent_company_code == ROOT

def test_relationship_hash_covers_both_report_scopes():
    from app.services.batch_project_workbook import entity_group_fields_hash
    standalone = SimpleNamespace(
        id=uuid4(), report_scope="standalone", parent_company_code=ROOT,
        ultimate_company_code=ROOT, relation_to_parent="subsidiary",
    )
    consolidated = SimpleNamespace(
        id=uuid4(), report_scope="consolidated", parent_company_code=ROOT,
        ultimate_company_code=ROOT, relation_to_parent="subsidiary",
    )
    before = entity_group_fields_hash([standalone, consolidated])
    consolidated.relation_to_parent = "branch"
    assert entity_group_fields_hash([standalone, consolidated]) != before

@pytest.mark.asyncio
async def test_http_validate_apply_and_status_use_one_persisted_plan(db):
    import httpx
    from fastapi import FastAPI
    from app.deps import get_current_user
    from app.core.database import get_db
    from app.models.base import UserRole
    from app.models.core import User
    from app.routers.batch_project import router

    user = User(
        username="batch_http_admin", email="batch_http_admin@example.test",
        hashed_password="x", role=UserRole.admin,
    )
    db.add(user)
    await db.commit()

    async def override_db():
        yield db

    app = FastAPI()
    app.include_router(router)
    app.dependency_overrides[get_db] = override_db
    app.dependency_overrides[get_current_user] = lambda: user
    transport = httpx.ASGITransport(app=app)
    key = "http-plan-idempotency-0001"
    async with httpx.AsyncClient(transport=transport, base_url="http://test") as client:
        validation_response = await client.post(
            "/api/projects/batch-validate?mode=project",
            files={"file": ("建项.xlsx", make_excel([row("HTTP建项", CHILD, 2034)]),
                            "application/vnd.openxmlformats-officedocument.spreadsheetml.sheet")},
            data={"idempotency_key": key},
        )
        assert validation_response.status_code == 200, validation_response.text
        validation = validation_response.json()
        assert validation["valid"] is True
        assert validation["validation_id"]
        assert "_plan_dependencies" not in validation["summary"]
        apply_response = await client.post("/api/projects/batch-apply", json={
            "validation_id": validation["validation_id"], "idempotency_key": key,
        })
        assert apply_response.status_code == 200, apply_response.text
        assert apply_response.json()["created_count"] == 1
        replay = await client.post("/api/projects/batch-apply", json={
            "validation_id": validation["validation_id"], "idempotency_key": key,
        })
        assert replay.status_code == 200
        assert replay.json()["idempotent"] is True
        status = await client.get(f"/api/projects/batch-imports/{validation['validation_id']}")
        assert status.status_code == 200
        assert status.json()["status"] == "applied"
        assert status.json()["mode"] == "project"
        assert status.json()["total_rows"] == 1
        assert status.json()["errors"] == []
        assert "_plan_dependencies" not in status.json()["summary"]

@pytest.mark.asyncio
async def test_empty_template_and_changed_header_are_blocked(db):
    empty = await generate_template()
    empty_result = await validate_batch(empty.getvalue(), db)
    assert empty_result.valid is False
    assert empty_result.errors[0].row_number == 0
    assert "没有可导入" in empty_result.errors[0].errors[0]

    workbook = Workbook()
    sheet = workbook.active
    sheet.title = "数据"
    bad_headers = list(_TEMPLATE_COLUMNS)
    bad_headers[1] = "企业编号"
    sheet.append(bad_headers)
    output = BytesIO()
    workbook.save(output)
    with pytest.raises(Exception) as captured:
        await validate_batch(output.getvalue(), db)
    assert getattr(captured.value, "status_code", None) == 400
    assert "表头" in str(getattr(captured.value, "detail", ""))


@pytest.mark.asyncio
async def test_invalid_plan_can_download_annotated_error_workbook(db):
    from app.services.batch_project_execution import build_plan_error_workbook
    actor = admin()
    plan, response, _ = await create_validation_plan(
        make_excel([["错误企业", "BAD", "简称", 2035, "年报审计", "企业会计准则", "单户"]]),
        db, actor=actor, file_name="错误.xlsx", mode="project",
        idempotency_key="invalid-plan-errors-0001",
    )
    assert response.valid is False
    output = await build_plan_error_workbook(db, plan_id=plan.id, actor=actor)
    workbook = load_workbook(output)
    rows = list(workbook["错误明细"].iter_rows(min_row=2, values_only=True))
    assert len(rows) == 1
    assert rows[0][-1]
    assert "企业代码" in rows[0][-1] or "18" in rows[0][-1]

@pytest.mark.asyncio
async def test_plan_apply_flush_failure_rolls_back_all_business_rows(engine, monkeypatch):
    factory = async_sessionmaker(engine, class_=AsyncSession, expire_on_commit=False)
    actor = admin()
    key = "atomic-failure-plan-0001"
    async with factory() as session:
        plan, validation, _ = await create_validation_plan(
            make_excel([
                row("原子甲", CHILD, 2036),
                row("原子乙", BRANCH, 2036),
            ]),
            session, actor=actor, file_name="原子.xlsx", mode="project",
            idempotency_key=key,
        )
        assert validation.valid is True
        plan_id = plan.id
        await session.commit()

    async with factory() as session:
        original_flush = session.flush
        calls = 0

        async def fail_business_flush(*args, **kwargs):
            nonlocal calls
            calls += 1
            if calls == 2:
                raise RuntimeError("injected apply flush failure")
            return await original_flush(*args, **kwargs)

        monkeypatch.setattr(session, "flush", fail_business_flush)
        with pytest.raises(RuntimeError, match="injected apply flush failure"):
            await apply_validation_plan(
                session, plan_id=plan_id, actor=actor, idempotency_key=key
            )
        await session.rollback()

    async with factory() as session:
        count = await session.scalar(select(func.count()).select_from(Project).where(
            Project.audit_year == 2036, Project.is_deleted == False,  # noqa: E712
        ))
        assert count == 0

@pytest.mark.asyncio
async def test_group_update_plan_rejects_relationship_changed_after_validation(engine):
    from app.models.audit_platform_schemas import BasicInfoSchema
    from app.services.batch_project_export import export_group_tree
    from app.services.project_wizard_service import create_project
    factory = async_sessionmaker(engine, class_=AsyncSession, expire_on_commit=False)
    async with factory() as session:
        root_project = await create_project(BasicInfoSchema(
            client_name="CAS集团", audit_year=2037, project_type="annual",
            accounting_standard="enterprise", company_code=ROOT,
            short_name="CAS集团", report_scope="consolidated",
            ultimate_company_code=ROOT,
        ), session)
        await create_project(BasicInfoSchema(
            client_name="CAS子公司", audit_year=2037, project_type="annual",
            accounting_standard="enterprise", company_code=CHILD,
            short_name="CAS子公司", report_scope="standalone",
            parent_company_code=ROOT, ultimate_company_code=ROOT,
            relation_to_parent="subsidiary",
        ), session)
        output, _ = await export_group_tree(root_project.id, session)
    workbook = load_workbook(output)
    child_row = next(cells for cells in workbook["数据"].iter_rows(min_row=2) if cells[1].value == CHILD)
    child_row[9].value = "分公司"
    edited = BytesIO()
    workbook.save(edited)
    actor = admin()
    key = "relationship-cas-plan-0001"
    async with factory() as session:
        plan, validation, _ = await create_validation_plan(
            edited.getvalue(), session, actor=actor, file_name="CAS.xlsx",
            mode="consolidation", idempotency_key=key,
        )
        assert validation.valid is True
        assert validation.summary.update == 1
        plan_id = plan.id
        await session.commit()
    async with factory() as session:
        child = (await session.execute(select(Project).where(
            Project.company_code == CHILD, Project.audit_year == 2037,
        ))).scalar_one()
        child.parent_company_code = None
        child.relation_to_parent = None
        await session.commit()
    async with factory() as session:
        with pytest.raises(Exception) as captured:
            await apply_validation_plan(
                session, plan_id=plan_id, actor=actor, idempotency_key=key
            )
        assert getattr(captured.value, "status_code", None) == 409
        assert "权限或库存已变化" in str(getattr(captured.value, "detail", ""))
        await session.rollback()

@pytest.mark.asyncio
async def test_plan_status_is_private_to_uploading_user(db):
    from app.services.batch_project_execution import get_owned_plan
    owner = admin()
    outsider = admin()
    plan, _, _ = await create_validation_plan(
        make_excel([row("私有计划", CHILD, 2038)]), db,
        actor=owner, file_name="私有.xlsx", mode="project",
        idempotency_key="private-plan-0001",
    )
    with pytest.raises(Exception) as captured:
        await get_owned_plan(db, plan.id, outsider)
    assert getattr(captured.value, "status_code", None) == 404

@pytest.mark.asyncio
async def test_plan_apply_batch_initializes_manager_memberships(db):
    from app.models.base import PermissionLevel, UserRole
    from app.models.core import ProjectUser, User
    from app.models.staff_models import ProjectAssignment, StaffMember
    manager = User(
        username="plan_manager", email="plan_manager@example.test",
        hashed_password="x", role=UserRole.manager,
    )
    db.add(manager)
    await db.flush()
    staff = StaffMember(user_id=manager.id, name="计划经理", employee_no="BP-MGR-2")
    db.add(staff)
    await db.commit()
    key = "manager-plan-batch-0001"
    plan, validation, _ = await create_validation_plan(
        make_excel([row("经理计划甲", CHILD, 2039), row("经理计划乙", BRANCH, 2039)]),
        db, actor=manager, file_name="经理.xlsx", mode="project", idempotency_key=key,
    )
    assert validation.valid is True
    await db.commit()
    result = await apply_validation_plan(
        db, plan_id=plan.id, actor=manager, idempotency_key=key
    )
    await db.commit()
    assert result.created_count == 2
    project_ids = [UUID(value) for value in result.project_ids]
    members = list((await db.execute(select(ProjectUser).where(
        ProjectUser.project_id.in_(project_ids),
        ProjectUser.user_id == manager.id,
        ProjectUser.is_deleted == False,  # noqa: E712
    ))).scalars().all())
    assignments = list((await db.execute(select(ProjectAssignment).where(
        ProjectAssignment.project_id.in_(project_ids),
        ProjectAssignment.staff_id == staff.id,
        ProjectAssignment.is_deleted == False,  # noqa: E712
    ))).scalars().all())
    assert len(members) == len(assignments) == 2
    assert {member.permission_level for member in members} == {PermissionLevel.edit}

def test_group_forest_nodes_expose_stable_node_key_for_frontend_expansion():
    from app.services.group_forest import build_forest
    project_id = uuid4()
    project = SimpleNamespace(
        id=project_id, company_code=CHILD, client_name="节点键公司",
        report_scope="standalone", audit_year=2040,
        audit_period_start=None, audit_period_end=None, wizard_state=None,
        name="节点键公司_2040", parent_company_code=None,
        ultimate_company_code=None, relation_to_parent=None,
        parent_project_id=None, status=None, consol_level=1,
    )
    forest = build_forest([project], year=2040)
    node = forest["independents"][0]
    assert node["nodeKey"] == f"{CHILD}@2040"

@pytest.mark.asyncio
async def test_unknown_template_version_is_rejected(db):
    output = await generate_template()
    workbook = load_workbook(output)
    workbook["_GT_META"]["B1"] = "999"
    changed = BytesIO()
    workbook.save(changed)
    with pytest.raises(Exception) as captured:
        await validate_batch(changed.getvalue(), db)
    assert getattr(captured.value, "status_code", None) == 400
    assert "模板版本 999 不受支持" in str(getattr(captured.value, "detail", ""))

@pytest.mark.asyncio
async def test_unauthorized_existing_entity_does_not_leak_group_fields(db):
    from app.models.audit_platform_schemas import BasicInfoSchema
    from app.models.base import UserRole
    from app.models.core import User
    from app.services.batch_project_plan import analyze_batch
    from app.services.batch_project_workbook import entity_group_fields_hash
    from app.services.project_wizard_service import create_project

    hidden = await create_project(BasicInfoSchema(
        client_name="隐藏企业", audit_year=2041, project_type="annual",
        accounting_standard="enterprise", company_code=CHILD,
        short_name="隐藏企业", report_scope="standalone",
        parent_company_code=ROOT, ultimate_company_code=ROOT,
        relation_to_parent="subsidiary",
    ), db)
    manager = User(
        username="hidden_probe_manager", email="hidden_probe@example.test",
        hashed_password="x", role=UserRole.manager,
    )
    db.add(manager)
    await db.commit()

    update_row = [
        "隐藏企业", CHILD, "隐藏企业", 2041, "年报审计", "企业会计准则", "单户",
        ROOT, ROOT, "分公司", "更新集团关系", str(hidden.id),
        entity_group_fields_hash([hidden]),
    ]
    update = await analyze_batch(
        make_excel([update_row]), db, current_user=manager, mode="consolidation"
    )
    messages = "；".join(message for item in update.response.errors for message in item.errors)
    assert "不存在或无权访问" in messages
    assert "项目ID与" not in messages
    assert "基线摘要" not in messages

    create_row = row("隐藏企业合并", CHILD, 2041, "合并")
    creation = await analyze_batch(
        make_excel([create_row]), db, current_user=manager, mode="consolidation"
    )
    absent = await analyze_batch(
        make_excel([row("不存在企业", BRANCH, 2041, "合并")]),
        db, current_user=manager, mode="consolidation",
    )
    # 非管理员预校验只以当前可见库存分类，不能用“隐藏但存在”作为无副作用的探测信号。
    assert creation.response.valid is absent.response.valid is True
    assert creation.response.errors == absent.response.errors == []
    assert creation.parsed_rows[0].parent_company_code == ""
    assert creation.parsed_rows[0].ultimate_company_code == ""
    # 兼容入口也不能把未显示的同年度库存误当作可新建项目并写入。
    legacy = await parse_and_import(
        make_excel([create_row]), db, current_user=manager, mode="consolidation"
    )
    assert legacy.success_count == 0
    assert legacy.fail_count == 1
    assert "权限或库存已变化" in legacy.failures[0].errors[0]

    persisted, persisted_validation, _ = await create_validation_plan(
        make_excel([row("隐藏企业单户", CHILD, 2041)]), db,
        actor=manager, file_name="隐藏冲突.xlsx", mode="project",
        idempotency_key="hidden-existing-plan-0001",
    )
    assert persisted_validation.valid is True
    await db.commit()
    with pytest.raises(Exception) as hidden_apply_error:
        await apply_validation_plan(
            db, plan_id=persisted.id, actor=manager,
            idempotency_key="hidden-existing-plan-0001",
        )
    assert getattr(hidden_apply_error.value, "status_code", None) == 409
    assert "权限或库存已变化" in str(getattr(hidden_apply_error.value, "detail", ""))
    await db.rollback()

@pytest.mark.asyncio
async def test_legacy_endpoint_never_treats_update_rows_as_new_projects(db):
    from app.models.audit_platform_schemas import BasicInfoSchema
    from app.services.batch_project_workbook import entity_group_fields_hash
    from app.services.project_wizard_service import create_project
    project = await create_project(BasicInfoSchema(
        client_name="兼容更新企业", audit_year=2042, project_type="annual",
        accounting_standard="enterprise", company_code=CHILD,
        short_name="兼容更新", report_scope="standalone",
    ), db)
    update_row = [
        project.client_name, project.company_code, project.short_name, 2042,
        "年报审计", "企业会计准则", "单户", "", "", "",
        "更新集团关系", str(project.id), entity_group_fields_hash([project]),
    ]
    result = await parse_and_import(make_excel([update_row]), db, mode="consolidation")
    assert result.success_count == 0
    assert result.fail_count == 1
    assert "确认执行" in result.failures[0].errors[0]
    count = await db.scalar(select(func.count()).select_from(Project).where(
        Project.company_code == CHILD, Project.audit_year == 2042,
        Project.is_deleted == False,  # noqa: E712
    ))
    assert count == 1

@pytest.mark.asyncio
async def test_expired_unapplied_plan_can_be_revalidated_with_same_key(db):
    from datetime import datetime, timedelta, timezone
    actor = admin()
    key = "expired-plan-revalidate-0001"
    payload = make_excel([row("过期重验", CHILD, 2043)])
    plan, first, _ = await create_validation_plan(
        payload, db, actor=actor, file_name="过期.xlsx", mode="project",
        idempotency_key=key,
    )
    assert first.valid is True
    plan.expires_at = datetime.now(timezone.utc) - timedelta(seconds=1)
    await db.commit()
    refreshed, second, reused = await create_validation_plan(
        payload, db, actor=actor, file_name="过期.xlsx", mode="project",
        idempotency_key=key,
    )
    assert reused is True
    assert refreshed.id == plan.id
    assert second.valid is True
    assert refreshed.expires_at > datetime.now(timezone.utc)

def test_exported_user_text_never_becomes_excel_formula():
    from app.services.batch_project_workbook import export_workbook
    prefixes = ["=1+1", "+SUM(1,1)", "-1+2", "@SUM(1,1)"]
    rows = [[value, CHILD, value, 2044, "年报审计", "企业会计准则", "单户",
             "", "", "", "新建", "", ""] for value in prefixes]
    output = export_workbook(rows, mode="project")
    sheet = load_workbook(output, data_only=False)["数据"]
    for row_number, value in enumerate(prefixes, start=2):
        assert sheet.cell(row_number, 1).value == value
        assert sheet.cell(row_number, 1).data_type == "s"
        assert sheet.cell(row_number, 3).data_type == "s"

@pytest.mark.asyncio
async def test_partner_creator_gets_project_access_without_fake_assignment(db):
    from app.deps import assert_project_permission
    from app.models.base import PermissionLevel, ProjectUserRole, UserRole
    from app.models.core import ProjectUser, User
    from app.models.staff_models import ProjectAssignment
    partner = User(
        username="batch_partner", email="batch_partner@example.test",
        hashed_password="x", role=UserRole.partner,
    )
    db.add(partner)
    await db.commit()
    key = "partner-plan-access-0001"
    plan, validation, _ = await create_validation_plan(
        make_excel([row("合伙人创建项目", CHILD, 2045)]), db,
        actor=partner, file_name="合伙人.xlsx", mode="project", idempotency_key=key,
    )
    assert validation.valid is True
    await db.commit()
    result = await apply_validation_plan(
        db, plan_id=plan.id, actor=partner, idempotency_key=key
    )
    await db.commit()
    project_id = UUID(result.project_ids[0])
    member = (await db.execute(select(ProjectUser).where(
        ProjectUser.project_id == project_id,
        ProjectUser.user_id == partner.id,
        ProjectUser.is_deleted == False,  # noqa: E712
    ))).scalar_one()
    assert member.role == ProjectUserRole.partner
    assert member.permission_level == PermissionLevel.edit
    assert await db.scalar(select(func.count()).select_from(ProjectAssignment).where(
        ProjectAssignment.project_id == project_id,
        ProjectAssignment.is_deleted == False,  # noqa: E712
    )) == 0
    await assert_project_permission(db, partner, project_id, "edit")

@pytest.mark.asyncio
async def test_existing_counterpart_conflict_is_repaired_not_marked_unchanged(engine):
    from app.models.audit_platform_schemas import BasicInfoSchema
    from app.services.batch_project_export import export_group_tree
    from app.services.project_wizard_service import create_project
    factory = async_sessionmaker(engine, class_=AsyncSession, expire_on_commit=False)
    async with factory() as session:
        root_project = await create_project(BasicInfoSchema(
            client_name="冲突集团", audit_year=2046, project_type="annual",
            accounting_standard="enterprise", company_code=ROOT,
            short_name="冲突集团", report_scope="consolidated", ultimate_company_code=ROOT,
        ), session)
        standalone = await create_project(BasicInfoSchema(
            client_name="冲突子公司", audit_year=2046, project_type="annual",
            accounting_standard="enterprise", company_code=CHILD,
            short_name="冲突单户", report_scope="standalone",
            parent_company_code=ROOT, ultimate_company_code=ROOT,
            relation_to_parent="subsidiary",
        ), session)
        consolidated = await create_project(BasicInfoSchema(
            client_name="冲突子公司", audit_year=2046, project_type="annual",
            accounting_standard="enterprise", company_code=CHILD,
            short_name="冲突合并", report_scope="consolidated",
            parent_company_code=ROOT, ultimate_company_code=ROOT,
            relation_to_parent="subsidiary",
        ), session)
        consolidated.relation_to_parent = "branch"
        await session.commit()
        output, _ = await export_group_tree(root_project.id, session)

    actor = admin()
    key = "counterpart-conflict-plan-0001"
    async with factory() as session:
        plan, validation, _ = await create_validation_plan(
            output.getvalue(), session, actor=actor, file_name="冲突树.xlsx",
            mode="consolidation", idempotency_key=key,
        )
        assert validation.valid is True, validation.errors
        assert validation.summary.update == 1
        plan_id = plan.id
        await session.commit()
    async with factory() as session:
        result = await apply_validation_plan(
            session, plan_id=plan_id, actor=actor, idempotency_key=key
        )
        await session.commit()
        assert result.updated_count == 1
        rows = list((await session.execute(select(Project).where(
            Project.company_code == CHILD, Project.audit_year == 2046,
        ))).scalars().all())
        assert {item.id for item in rows} == {standalone.id, consolidated.id}
        assert {item.relation_to_parent for item in rows} == {"subsidiary"}


@pytest.mark.asyncio
async def test_revoked_group_dependency_hides_plan_and_blocks_create_apply(db):
    from app.models.audit_platform_schemas import BasicInfoSchema
    from app.models.base import PermissionLevel, ProjectUserRole, UserRole
    from app.models.core import ProjectUser, User
    from app.services.batch_project_execution import get_readable_plan
    from app.services.project_wizard_service import create_project

    manager = User(
        username="revoked_batch_manager", email="revoked_batch@example.test",
        hashed_password="x", role=UserRole.manager,
    )
    db.add(manager)
    await db.flush()
    root = await create_project(BasicInfoSchema(
        client_name="撤权集团", audit_year=2047, project_type="annual",
        accounting_standard="enterprise", company_code=ROOT,
        short_name="撤权集团", report_scope="consolidated", ultimate_company_code=ROOT,
    ), db)
    membership = ProjectUser(
        project_id=root.id, user_id=manager.id, role=ProjectUserRole.manager,
        permission_level=PermissionLevel.edit,
    )
    db.add(membership)
    await db.commit()

    key = "revoked-plan-dependency-0001"
    plan, validation, _ = await create_validation_plan(
        make_excel([row("撤权后不得接入", BRANCH, 2047, parent=ROOT, ultimate=ROOT, relation="子公司")]),
        db, actor=manager, file_name="撤权.xlsx", mode="project", idempotency_key=key,
    )
    assert validation.valid is True
    await db.commit()

    membership.is_deleted = True
    await db.commit()
    with pytest.raises(Exception) as status_error:
        await get_readable_plan(db, plan.id, manager)
    assert getattr(status_error.value, "status_code", None) == 404

    with pytest.raises(Exception) as apply_error:
        await apply_validation_plan(db, plan_id=plan.id, actor=manager, idempotency_key=key)
    assert getattr(apply_error.value, "status_code", None) == 409
    assert "权限或库存已变化" in str(getattr(apply_error.value, "detail", ""))
    await db.rollback()
    assert await db.scalar(select(func.count()).select_from(Project).where(
        Project.company_code == BRANCH, Project.audit_year == 2047,
        Project.is_deleted == False,  # noqa: E712
    )) == 0


@pytest.mark.asyncio
async def test_create_plan_rejects_root_appearing_after_preview(db):
    from app.models.audit_platform_schemas import BasicInfoSchema
    from app.services.project_wizard_service import create_project

    actor = admin()
    key = "root-inventory-cas-0001"
    plan, validation, _ = await create_validation_plan(
        make_excel([row("库存 CAS 子公司", CHILD, 2048, parent=ROOT, ultimate=ROOT, relation="子公司")]),
        db, actor=actor, file_name="库存 CAS.xlsx", mode="project", idempotency_key=key,
    )
    assert validation.valid is True
    assert validation.summary.auto_root == 1
    await db.commit()

    await create_project(BasicInfoSchema(
        client_name="确认后出现的集团", audit_year=2048, project_type="annual",
        accounting_standard="enterprise", company_code=ROOT,
        short_name="后出现集团", report_scope="consolidated", ultimate_company_code=ROOT,
    ), db)
    await db.commit()

    with pytest.raises(Exception) as captured:
        await apply_validation_plan(db, plan_id=plan.id, actor=actor, idempotency_key=key)
    assert getattr(captured.value, "status_code", None) == 409
    assert "权限或库存已变化" in str(getattr(captured.value, "detail", ""))
    await db.rollback()
    assert await db.scalar(select(func.count()).select_from(Project).where(
        Project.company_code == CHILD, Project.audit_year == 2048,
        Project.is_deleted == False,  # noqa: E712
    )) == 0


@pytest.mark.asyncio
async def test_template_metadata_binds_mode_and_group_root_scope(db):
    from app.models.audit_platform_schemas import BasicInfoSchema
    from app.services.batch_project_export import _project_export_row
    from app.services.batch_project_plan import analyze_batch
    from app.services.batch_project_workbook import entity_group_fields_hash, export_workbook
    from app.services.project_wizard_service import create_project

    project_template = await generate_template(mode="project")
    with pytest.raises(Exception) as mode_error:
        await validate_batch(project_template.getvalue(), db, mode="consolidation")
    assert getattr(mode_error.value, "status_code", None) == 400
    assert "场景" in str(getattr(mode_error.value, "detail", ""))

    root = await create_project(BasicInfoSchema(
        client_name="来源集团", audit_year=2049, project_type="annual",
        accounting_standard="enterprise", company_code=ROOT,
        short_name="来源集团", report_scope="consolidated", ultimate_company_code=ROOT,
    ), db)
    outsider = await create_project(BasicInfoSchema(
        client_name="其他集团企业", audit_year=2049, project_type="annual",
        accounting_standard="enterprise", company_code=PARENT,
        short_name="其他企业", report_scope="standalone", ultimate_company_code=PARENT,
    ), db)
    await db.commit()
    scoped = export_workbook(
        [_project_export_row(outsider, action="update", base_hash=entity_group_fields_hash([outsider]))],
        mode="consolidation", root_project_id=str(root.id),
    )
    with pytest.raises(Exception) as scope_error:
        await analyze_batch(scoped.getvalue(), db, current_user=admin(), mode="consolidation")
    assert getattr(scope_error.value, "status_code", None) == 400
    assert "不属于模板来源集团" in str(getattr(scope_error.value, "detail", ""))

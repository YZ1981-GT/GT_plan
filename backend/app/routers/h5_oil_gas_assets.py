"""H5 油气资产 — 导入导出 + 行业检查端点

4个端点（POST/GET）：
- POST /api/h5/export-template   导出空白xlsx模板
- POST /api/h5/export-data       导出当前数据xlsx
- POST /api/h5/import-data       导入xlsx（multipart/form-data）
- GET  /api/h5/industry-check    检查项目行业适用性

科目编码: 1631油气资产 + 1632累计折耗（资产类/备抵类）
行业限制: oil_gas / mining
前端匹配: useH5ImportExport.ts
RFC5987 Content-Disposition header with Chinese filename encoding

Requirements: 3.3, 12.3
"""

from __future__ import annotations

import io
import json
import logging
from typing import Any
from urllib.parse import quote
from uuid import UUID

import sqlalchemy as sa
from fastapi import APIRouter, Depends, File, Form, HTTPException, UploadFile
from fastapi.responses import StreamingResponse
from openpyxl import Workbook, load_workbook
from openpyxl.styles import Alignment, Font, PatternFill
from openpyxl.utils import get_column_letter
from pydantic import BaseModel, Field
from sqlalchemy.ext.asyncio import AsyncSession

from app.core.database import get_db
from app.deps import get_current_user
from app.models.core import User

logger = logging.getLogger(__name__)

router = APIRouter(
    prefix="/api/h5",
    tags=["H5 油气资产"],
)

# ═══════════════════════════════════════════════════════════════════════════════
# Constants
# ═══════════════════════════════════════════════════════════════════════════════

_ROW_LIMIT = 500
_APPLICABLE_INDUSTRIES = {"oil_gas", "mining"}

# 支持导入导出的sheet
_SUPPORTED_SHEETS: set[str] = {"H5-2", "H5-4", "H5-7", "H5-8", "H5-16", "H5-17", "H5-18", "H5-19"}

# 各sheet列头定义
_SHEET_HEADERS: dict[str, list[str]] = {
    "H5-2": ["资产分类", "资产名称", "油田", "区块", "期初原值", "本期增加", "本期减少",
             "期初折耗", "本期计提", "本期转回", "备注"],
    "H5-4": ["序号", "资产名称", "油田", "原值", "净值", "闲置原因", "闲置起始日",
             "处置计划", "减值迹象", "核查结论", "备注"],
    "H5-7": ["序号", "资产名称", "油田", "区块", "金额", "勘探阶段", "资本化依据",
             "审批文件", "凭证号", "是否资本化", "结论"],
    "H5-8": ["序号", "资产名称", "油田", "处置原因", "原值", "累计折耗", "净值",
             "处置收入", "审批文件", "凭证号", "联动H10", "结论"],
    "H5-16": ["序号", "资产名称", "油田", "证照号", "类型", "有效期起", "有效期止",
              "权属主体", "抵押状态", "抵押金额", "结论"],
    "H5-17": ["序号", "关联方", "关联关系", "交易类型", "资产名称", "交易金额",
              "交易价格", "市场价格", "定价依据", "结论"],
    "H5-18": ["序号", "资产名称", "油田", "承租方", "租期(月)", "年租金", "资产净值",
              "合同号", "结论"],
    "H5-19": ["序号", "资产名称", "油田", "承租方", "本金", "未确认收益", "内含利率(%)",
              "本期利息", "剩余本金", "合同号", "结论"],
}

# 中文列名 → JSON 字段名映射
_FIELD_MAPS: dict[str, dict[str, str]] = {
    "H5-2": {
        "资产分类": "category", "资产名称": "name", "油田": "oilField", "区块": "block",
        "期初原值": "originalCostBegin", "本期增加": "originalCostIncrease",
        "本期减少": "originalCostDecrease", "期初折耗": "accDepletionBegin",
        "本期计提": "accDepletionProvision", "本期转回": "accDepletionReversal", "备注": "remark",
    },
    "H5-4": {
        "序号": "seq", "资产名称": "assetName", "油田": "oilField", "原值": "originalCost",
        "净值": "netValue", "闲置原因": "idleReason", "闲置起始日": "idleStartDate",
        "处置计划": "disposalPlan", "减值迹象": "impairmentSign", "核查结论": "conclusion", "备注": "remark",
    },
    "H5-7": {
        "序号": "seq", "资产名称": "assetName", "油田": "oilField", "区块": "block",
        "金额": "amount", "勘探阶段": "explorationStage", "资本化依据": "capitalizeReason",
        "审批文件": "approvalDoc", "凭证号": "voucherNo", "是否资本化": "isCapitalized", "结论": "conclusion",
    },
    "H5-8": {
        "序号": "seq", "资产名称": "assetName", "油田": "oilField", "处置原因": "disposalReason",
        "原值": "originalCost", "累计折耗": "accDepletion", "净值": "netValue",
        "处置收入": "disposalIncome", "审批文件": "approvalDoc", "凭证号": "voucherNo",
        "联动H10": "linkedH10", "结论": "conclusion",
    },
    "H5-16": {
        "序号": "seq", "资产名称": "assetName", "油田": "oilField", "证照号": "miningLicenseNo",
        "类型": "licenseType", "有效期起": "validFrom", "有效期止": "validTo",
        "权属主体": "ownerEntity", "抵押状态": "pledgeStatus", "抵押金额": "pledgeAmount", "结论": "conclusion",
    },
    "H5-17": {
        "序号": "seq", "关联方": "relatedParty", "关联关系": "relationship",
        "交易类型": "transType", "资产名称": "assetName", "交易金额": "transAmount",
        "交易价格": "transPrice", "市场价格": "marketPrice", "定价依据": "pricingBasis", "结论": "conclusion",
    },
    "H5-18": {
        "序号": "seq", "资产名称": "assetName", "油田": "oilField", "承租方": "lessee",
        "租期(月)": "leaseTerm", "年租金": "annualRent", "资产净值": "netValue",
        "合同号": "contractNo", "结论": "conclusion",
    },
    "H5-19": {
        "序号": "seq", "资产名称": "assetName", "油田": "oilField", "承租方": "lessee",
        "本金": "principal", "未确认收益": "unrecognizedFinanceIncome",
        "内含利率(%)": "interestRate", "本期利息": "currentInterest",
        "剩余本金": "remainingPrincipal", "合同号": "contractNo", "结论": "conclusion",
    },
}

# 数值类型字段
_NUMERIC_FIELDS: dict[str, set[str]] = {
    "H5-2": {"originalCostBegin", "originalCostIncrease", "originalCostDecrease",
             "accDepletionBegin", "accDepletionProvision", "accDepletionReversal"},
    "H5-4": {"seq", "originalCost", "netValue"},
    "H5-7": {"seq", "amount"},
    "H5-8": {"seq", "originalCost", "accDepletion", "netValue", "disposalIncome"},
    "H5-16": {"seq", "pledgeAmount"},
    "H5-17": {"seq", "transAmount", "transPrice", "marketPrice"},
    "H5-18": {"seq", "leaseTerm", "annualRent", "netValue"},
    "H5-19": {"seq", "principal", "unrecognizedFinanceIncome", "interestRate", "currentInterest", "remainingPrincipal"},
}

# checklist_responses 存储 item_id
_SHEET_ITEM_ID: dict[str, str] = {
    "H5-2": "H5-2-rows",
    "H5-4": "H5-4-rows",
    "H5-7": "H5-7-rows",
    "H5-8": "H5-8-rows",
    "H5-16": "H5-16-rows",
    "H5-17": "H5-17-rows",
    "H5-18": "H5-18-rows",
    "H5-19": "H5-19-rows",
}


# ═══════════════════════════════════════════════════════════════════════════════
# Pydantic Models
# ═══════════════════════════════════════════════════════════════════════════════


class H5ExportRequest(BaseModel):
    """导出请求体"""
    wpId: str = Field(..., description="底稿ID")
    sheet: str = Field(default="H5-2", description="目标sheet编码")


# ═══════════════════════════════════════════════════════════════════════════════
# GET /api/h5/industry-check — 行业适用性检查
# ═══════════════════════════════════════════════════════════════════════════════


@router.get("/industry-check")
async def h5_industry_check(
    project_id: str,
    db: AsyncSession = Depends(get_db),
    user: User = Depends(get_current_user),
) -> dict[str, Any]:
    """检查项目行业是否适用H5油气资产底稿"""
    from app.services.h5_oil_gas_assets_service import check_industry_applicability

    result = await check_industry_applicability(db, project_id)
    return result


# ═══════════════════════════════════════════════════════════════════════════════
# POST /api/h5/export-template — 导出空白模板
# ═══════════════════════════════════════════════════════════════════════════════


@router.post("/export-template")
async def h5_export_template(
    body: H5ExportRequest,
    user: User = Depends(get_current_user),
) -> StreamingResponse:
    """导出H5空白xlsx模板"""
    sheet = body.sheet
    if sheet not in _SUPPORTED_SHEETS:
        raise HTTPException(status_code=400, detail=f"不支持的sheet: {sheet}")

    headers = _SHEET_HEADERS[sheet]
    wb = Workbook()
    ws = wb.active
    ws.title = sheet

    # 标题行样式
    header_font = Font(bold=True, size=11)
    header_fill = PatternFill(start_color="D9E2F3", end_color="D9E2F3", fill_type="solid")
    header_align = Alignment(horizontal="center", vertical="center", wrap_text=True)

    for col_idx, col_name in enumerate(headers, 1):
        cell = ws.cell(row=1, column=col_idx, value=col_name)
        cell.font = header_font
        cell.fill = header_fill
        cell.alignment = header_align
        ws.column_dimensions[get_column_letter(col_idx)].width = max(12, len(col_name) * 2 + 4)

    buf = io.BytesIO()
    wb.save(buf)
    buf.seek(0)

    filename = f"H5油气资产_{sheet}_模板.xlsx"
    encoded = quote(filename)
    return StreamingResponse(
        buf,
        media_type="application/vnd.openxmlformats-officedocument.spreadsheetml.sheet",
        headers={
            "Content-Disposition": f"attachment; filename*=UTF-8''{encoded}",
        },
    )


# ═══════════════════════════════════════════════════════════════════════════════
# POST /api/h5/export-data — 导出当前数据
# ═══════════════════════════════════════════════════════════════════════════════


@router.post("/export-data")
async def h5_export_data(
    body: H5ExportRequest,
    db: AsyncSession = Depends(get_db),
    user: User = Depends(get_current_user),
) -> StreamingResponse:
    """导出H5当前数据xlsx"""
    sheet = body.sheet
    if sheet not in _SUPPORTED_SHEETS:
        raise HTTPException(status_code=400, detail=f"不支持的sheet: {sheet}")

    # 从 checklist_responses 读取数据
    item_id = _SHEET_ITEM_ID[sheet]
    rows_data = await _load_rows_from_db(db, body.wpId, item_id)

    headers = _SHEET_HEADERS[sheet]
    field_map = _FIELD_MAPS[sheet]
    # 反转映射：field_name → column_name
    reverse_map = {v: k for k, v in field_map.items()}

    wb = Workbook()
    ws = wb.active
    ws.title = sheet

    # 写入表头
    header_font = Font(bold=True, size=11)
    header_fill = PatternFill(start_color="D9E2F3", end_color="D9E2F3", fill_type="solid")
    for col_idx, col_name in enumerate(headers, 1):
        cell = ws.cell(row=1, column=col_idx, value=col_name)
        cell.font = header_font
        cell.fill = header_fill
        ws.column_dimensions[get_column_letter(col_idx)].width = max(12, len(col_name) * 2 + 4)

    # 写入数据行
    for row_idx, row_data in enumerate(rows_data, 2):
        for col_idx, col_name in enumerate(headers, 1):
            field = field_map.get(col_name, "")
            value = row_data.get(field, "")
            # 布尔值转文本
            if isinstance(value, bool):
                value = "是" if value else "否"
            ws.cell(row=row_idx, column=col_idx, value=value)

    buf = io.BytesIO()
    wb.save(buf)
    buf.seek(0)

    filename = f"H5油气资产_{sheet}_数据.xlsx"
    encoded = quote(filename)
    return StreamingResponse(
        buf,
        media_type="application/vnd.openxmlformats-officedocument.spreadsheetml.sheet",
        headers={
            "Content-Disposition": f"attachment; filename*=UTF-8''{encoded}",
        },
    )


# ═══════════════════════════════════════════════════════════════════════════════
# POST /api/h5/import-data — 导入xlsx
# ═══════════════════════════════════════════════════════════════════════════════


@router.post("/import-data")
async def h5_import_data(
    wp_id: str = Form(...),
    sheet: str = Form(default="H5-2"),
    file: UploadFile = File(...),
    db: AsyncSession = Depends(get_db),
    user: User = Depends(get_current_user),
) -> dict[str, Any]:
    """导入xlsx数据到H5底稿"""
    if sheet not in _SUPPORTED_SHEETS:
        raise HTTPException(status_code=400, detail=f"不支持的sheet: {sheet}")

    # 读取xlsx
    content = await file.read()
    try:
        wb = load_workbook(io.BytesIO(content), read_only=True, data_only=True)
    except Exception as e:
        raise HTTPException(status_code=400, detail=f"无法解析xlsx文件: {e}")

    ws = wb.active
    if ws is None:
        raise HTTPException(status_code=400, detail="xlsx文件无有效sheet")

    # 解析表头
    headers_row = [str(cell.value or "").strip() for cell in ws[1]]
    field_map = _FIELD_MAPS[sheet]
    numeric_fields = _NUMERIC_FIELDS.get(sheet, set())

    rows: list[dict[str, Any]] = []
    for row_idx, row in enumerate(ws.iter_rows(min_row=2, values_only=True), 1):
        if row_idx > _ROW_LIMIT:
            break
        if all(cell is None or str(cell).strip() == "" for cell in row):
            continue

        row_dict: dict[str, Any] = {
            "rowId": f"imported-{row_idx}",
        }
        for col_idx, col_name in enumerate(headers_row):
            field = field_map.get(col_name)
            if not field:
                continue
            val = row[col_idx] if col_idx < len(row) else None

            if field in numeric_fields:
                try:
                    row_dict[field] = float(val) if val is not None else 0
                except (ValueError, TypeError):
                    row_dict[field] = 0
            elif field == "impairmentSign" or field == "isCapitalized" or field == "linkedH10":
                row_dict[field] = str(val).strip() in ("是", "True", "true", "1", "TRUE")
            else:
                row_dict[field] = str(val).strip() if val is not None else ""

        rows.append(row_dict)

    wb.close()

    # 存入 checklist_responses
    item_id = _SHEET_ITEM_ID[sheet]
    await _save_rows_to_db(db, wp_id, item_id, rows)

    return {
        "imported_count": len(rows),
        "sheet": sheet,
        "message": f"成功导入 {len(rows)} 行数据到 {sheet}",
    }


# ═══════════════════════════════════════════════════════════════════════════════
# DB helpers
# ═══════════════════════════════════════════════════════════════════════════════


async def _load_rows_from_db(db: AsyncSession, wp_id: str, item_id: str) -> list[dict]:
    """从 checklist_responses 加载行数据"""
    result = await db.execute(
        sa.text("""
            SELECT remark FROM checklist_responses
            WHERE wp_id = :wp_id AND item_id = :item_id
            LIMIT 1
        """),
        {"wp_id": wp_id, "item_id": item_id},
    )
    row = result.first()
    if not row or not row[0]:
        return []
    try:
        data = json.loads(row[0])
        return data if isinstance(data, list) else []
    except (json.JSONDecodeError, TypeError):
        return []


async def _save_rows_to_db(db: AsyncSession, wp_id: str, item_id: str, rows: list[dict]) -> None:
    """保存行数据到 checklist_responses（upsert）"""
    json_data = json.dumps(rows, ensure_ascii=False)
    await db.execute(
        sa.text("""
            INSERT INTO checklist_responses (id, wp_id, item_id, remark, conclusion)
            VALUES (gen_random_uuid(), :wp_id, :item_id, :remark, NULL)
            ON CONFLICT (wp_id, item_id)
            DO UPDATE SET remark = :remark
        """),
        {"wp_id": wp_id, "item_id": item_id, "remark": json_data},
    )
    await db.flush()

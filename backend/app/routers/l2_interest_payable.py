"""L2 应付利息 — 导入导出三级端点 + 计提核对API

4个端点：
- GET  /api/l2-interest-payable/{wp_id}/export-template  空白模板xlsx
- GET  /api/l2-interest-payable/{wp_id}/export-data      数据xlsx
- POST /api/l2-interest-payable/{wp_id}/import-data      解析xlsx写入
- GET  /api/l2-interest-payable/{wp_id}/accrual-check    计提核对（L1/L3 vs L2）

科目编码: 2231应付利息（贷方/负债类）
RFC5987 Content-Disposition header with Chinese filename encoding

Requirements: 6.5
"""

from __future__ import annotations

import inspect
import io
import json
import logging
import re
from typing import Any
from urllib.parse import quote
from uuid import uuid4

from fastapi import APIRouter, Depends, File, HTTPException, Query, UploadFile
from fastapi.responses import StreamingResponse
from openpyxl import Workbook, load_workbook
from openpyxl.styles import Alignment, Font, PatternFill
from sqlalchemy.ext.asyncio import AsyncSession

from app.core.database import get_db
from app.deps import get_current_user
from app.models.core import User

from app.routers.wp_render_strategies._x3_adjustment_import_export import (
    X3_SHEET_SPECS,
    X3ContractError,
    attach_shape_a_routes,
)

logger = logging.getLogger(__name__)

router = APIRouter(prefix="/api/l2-interest-payable", tags=["l2-interest-payable"])

# ═══════════════════════════════════════════════════════════════════════════════
# Sheet 配置
# ═══════════════════════════════════════════════════════════════════════════════

_ROW_LIMIT = 500

_SHEET_NAME = "L2-2 明细表"

_HEADERS: list[str] = [
    "来源类别", "合同名称", "币种", "本金", "年利率(%)",
    "计息起", "计息止", "期初应付", "本期计提", "本期支付",
]

# 中文列名 → JSON 字段名
_FIELD_MAP: dict[str, str] = {
    "来源类别": "sourceCategory",
    "合同名称": "contractName",
    "币种": "currency",
    "本金": "principal",
    "年利率(%)": "annualRate",
    "计息起": "interestStartDate",
    "计息止": "interestEndDate",
    "期初应付": "openingPayable",
    "本期计提": "currentAccrual",
    "本期支付": "currentPayment",
}

# 数值字段
_NUMERIC_FIELDS: set[str] = {
    "principal", "annualRate", "openingPayable", "currentAccrual", "currentPayment",
}

# checklist_responses item_id
_ITEM_ID = "L2-detail-rows"


# ═══════════════════════════════════════════════════════════════════════════════
# X-3 调整分录汇总表 —— 形态 A 三态 + 形态 B 的 `sheet` 分派
# spec: x3-adjustment-entry-import-export 任务 5.2（R4.1 / R4.2 / R4.3 / R4.4 / R1.6）
# ═══════════════════════════════════════════════════════════════════════════════

#: 本模块的短前缀 = catalog `import_export.api_prefix`，单份 UI 与批量两条通路共用它。
_IE_API_PREFIX = "l2"

#: 本前缀的 X-3 sheet 码 —— 从 `X3_SHEET_SPECS`（Key_Ledger 单一真源）按短前缀**反查**。
#: 🔴 本模块内不写 X-3 字面量：写死一份即成后端第二真源，前端改键时不会打红。
#: 反查为空 ⇒ 下面 `attach_shape_a_routes` 立即抛 `X3ContractError`（fail-loud，不静默少端点）。
_X3_CODES: frozenset[str] = frozenset(
    code for code, spec in X3_SHEET_SPECS.items() if spec.api_prefix == _IE_API_PREFIX
)

#: 形态 B（长前缀）三态端点服务的业务表码 —— 从既有 `_SHEET_NAME` **派生**（首段即 sheet 码），
#: 故 sheet 码在本模块只写一次、结构上不可能与既有导出的 `ws.title` 分叉。
#:
#: 🔴 派生失败必须**立即抛**而不是留一个错的白名单：前端 `useL2ImportExport` 一直在传该值
#:    （`L2TabDetail.vue` 三处调用各传业务表码），白名单错了会让线上导入导出全落进下面的
#:    400 分支 —— 接口红但四层守卫全绿，属最贵的一类。
_LEGACY_SHEET_CODE = _SHEET_NAME.split(" ", 1)[0]
if not re.fullmatch(rf"{_IE_API_PREFIX.upper()}-\d+", _LEGACY_SHEET_CODE):
    raise X3ContractError(
        f"{_IE_API_PREFIX} 既有 sheet 码派生失败: _SHEET_NAME={_SHEET_NAME!r} "
        f"得到 {_LEGACY_SHEET_CODE!r}，期望 {_IE_API_PREFIX.upper()}-N 形态"
    )

#: 形态 B 可服务的 sheet 集合。取值与施加前逐字相同（既有 handler 只构建这一张表）。
_LEGACY_SHEETS: frozenset[str] = frozenset({_LEGACY_SHEET_CODE})

#: 本前缀 sheet 白名单的**唯一真源**（R4.4）= 形态 B 的业务表 + 形态 A 的 X-3。
#: 🔴 X-3 刻意不进 `_LEGACY_SHEETS`：既有 handler 的表头/字段/`item_id` 都只有业务表那一套，
#:    放进去会让 X-3 走到取不到配置的旧路径（串表导出）。X-3 只由共享实现服务。
IE_SHEETS: frozenset[str] = _LEGACY_SHEETS | _X3_CODES

attach_shape_a_routes(router, api_prefix=_IE_API_PREFIX, sheets=_X3_CODES)

#: 形态 A 三态端点的可调用对象（挂载后按路径后缀反查）。
#: 🔴 委派必须落到**同一个** callable 上，不抄一份编排（10MB 上限 / 解析失败转 400 /
#:    残留清理策略 / 响应载荷）：抄一份就是 R1.5 禁的第三套实现，两侧此后各自漂移。
_X3_SHAPE_A: dict[str, Any] = {
    str(getattr(route, "path", "")).rsplit("/", 1)[-1]: route.endpoint
    for route in router.routes
    if str(getattr(route, "path", "")).startswith(
        f"/api/workpapers/{{wp_id}}/{_IE_API_PREFIX}/"
    )
}
if len(_X3_SHAPE_A) != 3:
    raise X3ContractError(
        f"{_IE_API_PREFIX} 形态 A 三态端点反查到 {sorted(_X3_SHAPE_A)}（期望 3 条）"
    )


def _requires_x3(sheet: str | None) -> bool:
    """形态 B handler 的 `sheet` 三分支判定（R4.1 / R4.2 / R1.6）。

    - `None` ⇒ False：既有调用方不传该参数，走**逐字不变**的既有路径；
    - 值 ∈ `IE_SHEETS` 且为 X-3 ⇒ True：委派共享实现；
    - 值 ∈ `IE_SHEETS` 的其余业务表 ⇒ False：既有路径，产物与施加前逐字节相同；
    - 值 ∉ `IE_SHEETS` ⇒ 400（文案沿用共享实现 `_validate_sheet`）。

    🔴 未登记值刻意**不**回退成「导出全部 sheet」：那是 E6 存量缺陷的形态，
       200 + 错内容比 400 难查一个量级（父 spec 在 L4-3 上踩过）。
    """
    if sheet is None:
        return False
    if sheet not in IE_SHEETS:
        raise HTTPException(400, f"不支持的sheet: {sheet}。支持: {sorted(IE_SHEETS)}")
    return sheet in _X3_CODES


async def _x3_delegate(action: str, **kwargs: Any) -> Any:
    """把请求交给形态 A 的同名端点（= 共享实现的那个 callable）。

    调用方没给的**非依赖**参数从共享实现签名的默认值剥出（`Query(...)` 包装 → `.default`）：
    🔴 不剥而直接调用，`strategy` 会拿到 `Query` 包装对象本身 ⇒ 共享实现的
       `_should_purge_residual` 比对 `== "overwrite"` 不成立 ⇒ 残留族键不清 = 幽灵行，
       而接口仍 200。这类默认值只有共享实现一份，本模块不抄第二份。
    签名要求的参数本模块给不出时**抛 500**，不静默传个占位值。
    """
    endpoint = _X3_SHAPE_A[action]
    for name, par in inspect.signature(endpoint).parameters.items():
        if name in kwargs:
            continue
        raw = par.default
        if raw is inspect.Parameter.empty or getattr(raw, "dependency", None) is not None:
            raise HTTPException(
                500, f"X-3 委派缺参: {action} 的 {name!r} 须由本模块提供（共享实现签名已变）"
            )
        value = getattr(raw, "default", raw)
        if value is Ellipsis or type(value).__name__ == "PydanticUndefinedType":
            raise HTTPException(
                500, f"X-3 委派缺参: {action} 的 {name!r} 是必填项（共享实现签名已变）"
            )
        kwargs[name] = value
    return await endpoint(**kwargs)


# ═══════════════════════════════════════════════════════════════════════════════
# 辅助函数
# ═══════════════════════════════════════════════════════════════════════════════


def _safe_float(val: Any) -> float:
    if val is None:
        return 0.0
    try:
        return float(val)
    except (ValueError, TypeError):
        return 0.0


def _safe_str(val: Any) -> str:
    if val is None:
        return ""
    return str(val).strip()


def _create_template_wb() -> Workbook:
    """创建空白模板xlsx（含表头+格式，无数据行）"""
    wb = Workbook()
    ws = wb.active
    ws.title = _SHEET_NAME

    header_font = Font(bold=True, size=11)
    header_fill = PatternFill(start_color="D9E1F2", end_color="D9E1F2", fill_type="solid")
    header_align = Alignment(horizontal="center", vertical="center", wrap_text=True)

    for col_idx, col_name in enumerate(_HEADERS, 1):
        cell = ws.cell(row=1, column=col_idx, value=col_name)
        cell.font = header_font
        cell.fill = header_fill
        cell.alignment = header_align
        ws.column_dimensions[cell.column_letter].width = max(len(col_name) * 2 + 4, 12)

    ws.freeze_panes = "A2"
    return wb


def _validate_columns(ws: Any) -> list[str]:
    """校验列头，返回不匹配的列名列表"""
    expected = set(_HEADERS)
    actual: list[str] = []
    for cell in next(ws.iter_rows(min_row=1, max_row=1)):
        if cell.value is not None:
            actual.append(str(cell.value).strip())

    missing = [h for h in expected if h not in actual]
    return missing


def _get_actual_headers(ws: Any) -> list[str]:
    """获取worksheet实际列头"""
    headers: list[str] = []
    for cell in next(ws.iter_rows(min_row=1, max_row=1)):
        if cell.value is not None:
            headers.append(str(cell.value).strip())
    return headers


def _col_val(row: tuple, actual_headers: list[str], col_name: str) -> Any:
    """从行数据中按列名取值"""
    try:
        idx = actual_headers.index(col_name)
        return row[idx] if idx < len(row) else None
    except (ValueError, IndexError):
        return None


def _export_row(data: dict) -> list:
    """将JSON行数据按列头顺序转换为xlsx行"""
    result = []
    for col_name in _HEADERS:
        field = _FIELD_MAP.get(col_name, col_name)
        val = data.get(field, "")
        result.append(val if val is not None else "")
    return result


def _parse_row(row: tuple, actual_headers: list[str]) -> dict:
    """导入行解析——根据列名映射回JSON字段"""
    result: dict[str, Any] = {"rowId": str(uuid4())}
    for col_name, field_name in _FIELD_MAP.items():
        raw = _col_val(row, actual_headers, col_name)
        if field_name in _NUMERIC_FIELDS:
            result[field_name] = _safe_float(raw)
        else:
            result[field_name] = _safe_str(raw)
    return result


# ═══════════════════════════════════════════════════════════════════════════════
# 端点
# ═══════════════════════════════════════════════════════════════════════════════


@router.get("/{wp_id}/export-template")
async def l2_export_template(
    wp_id: str,
    sheet: str | None = Query(None),
    current_user: User = Depends(get_current_user),
) -> StreamingResponse:
    """导出空白xlsx模板（含表头+格式，无数据行）

    `sheet` 为 additive 可选参数（任务 5.2）：不传 ⇒ 行为与施加前逐字相同。

    Requirements: 6.5
    """
    if _requires_x3(sheet):
        return await _x3_delegate(
            "export-template", wp_id=wp_id, sheet=sheet, current_user=current_user
        )
    wb = _create_template_wb()
    buffer = io.BytesIO()
    wb.save(buffer)
    buffer.seek(0)

    filename = "L2应付利息_明细表_模板.xlsx"
    return StreamingResponse(
        buffer,
        media_type="application/vnd.openxmlformats-officedocument.spreadsheetml.sheet",
        headers={"Content-Disposition": f"attachment; filename*=UTF-8''{quote(filename)}"},
    )


@router.get("/{wp_id}/export-data")
async def l2_export_data(
    wp_id: str,
    sheet: str | None = Query(None),
    db: AsyncSession = Depends(get_db),
    current_user: User = Depends(get_current_user),
) -> StreamingResponse:
    """导出含当前数据的xlsx

    `sheet` 为 additive 可选参数（任务 5.2）：不传 ⇒ 行为与施加前逐字相同。

    Requirements: 6.5
    """
    if _requires_x3(sheet):
        return await _x3_delegate(
            "export-data", wp_id=wp_id, sheet=sheet, db=db, current_user=current_user
        )
    import sqlalchemy as sa

    rows_data: list[dict] = []

    result = await db.execute(
        sa.text(
            "SELECT remark FROM checklist_responses "
            "WHERE wp_id = :wp_id AND item_id = :item_id LIMIT 1"
        ),
        {"wp_id": wp_id, "item_id": _ITEM_ID},
    )
    row = result.fetchone()
    if row and row.remark:
        try:
            rows_data = json.loads(row.remark)
        except (json.JSONDecodeError, TypeError):
            pass

    # 生成 xlsx
    wb = _create_template_wb()
    ws = wb.active

    for data_row in rows_data:
        ws.append(_export_row(data_row))

    buffer = io.BytesIO()
    wb.save(buffer)
    buffer.seek(0)

    filename = "L2应付利息_明细表_数据.xlsx"
    return StreamingResponse(
        buffer,
        media_type="application/vnd.openxmlformats-officedocument.spreadsheetml.sheet",
        headers={"Content-Disposition": f"attachment; filename*=UTF-8''{quote(filename)}"},
    )


@router.post("/{wp_id}/import-data")
async def l2_import_data(
    wp_id: str,
    sheet: str | None = Query(None),
    file: UploadFile = File(...),
    db: AsyncSession = Depends(get_db),
    current_user: User = Depends(get_current_user),
) -> dict:
    """导入xlsx解析写入checklist_responses

    `sheet` 为 additive 可选参数（任务 5.2）：不传 ⇒ 行为与施加前逐字相同。

    Requirements: 6.5
    """
    if _requires_x3(sheet):
        return await _x3_delegate(
            "import-data",
            wp_id=wp_id,
            sheet=sheet,
            file=file,
            db=db,
            current_user=current_user,
        )
    import sqlalchemy as sa

    # 读取上传文件
    content = await file.read()
    if len(content) > 10 * 1024 * 1024:
        raise HTTPException(400, "文件大小超过 10MB 限制")

    try:
        wb = load_workbook(io.BytesIO(content), read_only=True, data_only=True)
    except Exception:
        raise HTTPException(400, "无法解析xlsx文件")

    ws = wb.active

    # 校验列头
    missing = _validate_columns(ws)
    if missing:
        raise HTTPException(400, f"列名不匹配，缺少: {missing}")

    actual_headers = _get_actual_headers(ws)

    # 解析数据行
    parsed_rows: list[dict] = []
    for row_idx, row in enumerate(ws.iter_rows(min_row=2, values_only=True), start=2):
        if row_idx > _ROW_LIMIT + 1:
            break
        # 跳过全空行
        if all(cell is None or str(cell).strip() == "" for cell in row):
            continue
        parsed_rows.append(_parse_row(row, actual_headers))

    # 写入 checklist_responses
    json_str = json.dumps(parsed_rows, ensure_ascii=False)

    await db.execute(
        sa.text(
            "INSERT INTO checklist_responses (id, wp_id, item_id, remark) "
            "VALUES (:id, :wp_id, :item_id, :remark) "
            "ON CONFLICT (wp_id, item_id) DO UPDATE SET remark = :remark"
        ),
        {
            "id": str(uuid4()),
            "wp_id": wp_id,
            "item_id": _ITEM_ID,
            "remark": json_str,
        },
    )
    await db.commit()

    return {"row_count": len(parsed_rows), "sheet": "L2-2"}


@router.get("/{wp_id}/accrual-check")
async def l2_accrual_check(
    wp_id: str,
    db: AsyncSession = Depends(get_db),
    current_user: User = Depends(get_current_user),
) -> dict:
    """获取计提核对数据（L1/L3利息测算 vs L2账面计提）

    从L1-5利息测算表和L2明细表读取数据，计算差异。
    L3长期借款利息暂用0（L3组件未就绪时）。

    Requirements: 6.5, 4.4
    """
    import sqlalchemy as sa

    # ─── 从L2明细表获取账面计提合计 ────────────────────────────────────────
    l2_booked: float = 0.0
    result = await db.execute(
        sa.text(
            "SELECT remark FROM checklist_responses "
            "WHERE wp_id = :wp_id AND item_id = :item_id LIMIT 1"
        ),
        {"wp_id": wp_id, "item_id": _ITEM_ID},
    )
    row = result.fetchone()
    if row and row.remark:
        try:
            detail_rows = json.loads(row.remark)
            for r in detail_rows:
                l2_booked += _safe_float(r.get("currentAccrual", 0))
        except (json.JSONDecodeError, TypeError):
            pass

    # ─── 从L1-5利息测算表获取L1利息合计 ───────────────────────────────────
    l1_interest: float = 0.0
    # L1利息测算存储在同project的L1底稿 checklist_responses，item_id='L1-interest-calc-rows'
    # 查找同project下的L1底稿wp_id
    try:
        l1_result = await db.execute(
            sa.text(
                "SELECT cr.remark FROM checklist_responses cr "
                "JOIN working_paper wp1 ON wp1.id = cr.wp_id "
                "JOIN working_paper wp2 ON wp2.project_id = wp1.project_id "
                "WHERE wp2.id = :wp_id "
                "AND cr.item_id = 'L1-interest-calc-rows' "
                "LIMIT 1"
            ),
            {"wp_id": wp_id},
        )
        l1_row = l1_result.fetchone()
        if l1_row and l1_row.remark:
            l1_rows = json.loads(l1_row.remark)
            for r in l1_rows:
                l1_interest += _safe_float(r.get("calculatedInterest", 0))
    except Exception as e:
        logger.warning("L2 accrual-check: L1利息查询失败: %s", e)

    # ─── L3长期借款利息（L3组件未就绪时返回0） ────────────────────────────
    l3_interest: float = 0.0
    try:
        l3_result = await db.execute(
            sa.text(
                "SELECT cr.remark FROM checklist_responses cr "
                "JOIN working_paper wp1 ON wp1.id = cr.wp_id "
                "JOIN working_paper wp2 ON wp2.project_id = wp1.project_id "
                "WHERE wp2.id = :wp_id "
                "AND cr.item_id = 'L3-interest-calc-rows' "
                "LIMIT 1"
            ),
            {"wp_id": wp_id},
        )
        l3_row = l3_result.fetchone()
        if l3_row and l3_row.remark:
            l3_rows = json.loads(l3_row.remark)
            for r in l3_rows:
                l3_interest += _safe_float(r.get("calculatedInterest", 0))
    except Exception as e:
        logger.warning("L2 accrual-check: L3利息查询失败: %s", e)

    # ─── 汇总 ────────────────────────────────────────────────────────────
    total_estimated = l1_interest + l3_interest
    diff = total_estimated - l2_booked
    is_consistent = abs(diff) < 0.01  # 差异小于1分视为一致

    return {
        "l1_interest": round(l1_interest, 2),
        "l3_interest": round(l3_interest, 2),
        "l2_booked": round(l2_booked, 2),
        "diff": round(diff, 2),
        "is_consistent": is_consistent,
    }

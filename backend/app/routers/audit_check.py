"""审计检查复核收口面板 — summary 端点（audit-check-review-gate-hardening Task 2.3）

`GET /api/projects/{pid}/audit-checks/summary`
  批量读项目所有底稿的统一检查项（`parsed_data.audit_checks`），无该字段时**退回**读
  legacy `parsed_data.fine_checks`（补 `source=fine_rule`）— 向后兼容（P13 / Req10.4）。
  每底稿返回新鲜度三态（`checked_at`/`updated_at`/`stale`/`never_checked`，P3 / Req1.2-1.4）；
  项目汇总 `ProjectCheckSummary`（通过率分母=已判定数不含 null，P1/P4 / Req5）。

设计要点（design §3/§4）：
- 一次批量 JOIN 查询取 `WorkingPaper.id`/`parsed_data`/`updated_at` + `WpIndex.wp_code`/
  `wp_name`/`audit_cycle`，避免 N+1。
- 新鲜度：`checked_at` = `audit_checks_at` 或 legacy `fine_extracted_at`（可空）；
  `stale` = `updated_at > checked_at`（`checked_at` 为空时 `stale=false`）；
  `never_checked` = `checked_at` 为空。时区/字符串比较稳健（统一转 aware datetime，
  解析失败按未过期处理，不抛）。
- 未检查底稿不贡献任何 check 项 → 不计入通过率（P4，其 checks 为空自然不计入）。
- 权限：项目只读（`require_project_access("readonly")`，P12 / Req10.1，不放宽）。

本端点是 `get_fine_checks_summary` 的升级版；旧端点 `GET /fine-checks/summary` 保留不动
（QC / 既有消费者）。
"""

from __future__ import annotations

import logging
from datetime import datetime, timezone
from uuid import UUID

import sqlalchemy as sa
from fastapi import APIRouter, Body, Depends, HTTPException
from pydantic import BaseModel
from sqlalchemy.ext.asyncio import AsyncSession

from app.core.database import get_db
from app.deps import require_project_access
from app.models.audit_platform_models import AuditCheckSignoff
from app.models.core import Project
from app.models.workpaper_models import WorkingPaper, WpIndex
from app.services.audit_check.aggregator import AuditCheckAggregator
from app.services.formula_management.delivery_export import (
    content_disposition_attachment,
)
from app.services.audit_check.models import (
    BACKEND_COMPUTED_SOURCES,
    FRONTEND_REPORTABLE_SOURCES,
    SEVERITY_INFO,
    AuditCheckItem,
    ProjectCheckSummary,
    from_fine_check_dict,
)
from sqlalchemy.orm.attributes import flag_modified

logger = logging.getLogger(__name__)

router = APIRouter(tags=["审计检查复核"])

# ═══════════════════════════════════════════
# 并发去重/幂等保护（Req2.3）
# ═══════════════════════════════════════════
# 进程内「重算进行中」项目集合。asyncio 单线程协作调度下，检查 + add 之间无 await
# → 原子；同项目重算进行中时重复触发直接返回 in_progress，不重复启动。
# 聚合器写缓存为覆盖写（天然幂等），即便并发也不产生脏数据累积。
_recompute_in_progress: set[str] = set()


# ═══════════════════════════════════════════
# 新鲜度工具（稳健时区/字符串比较，解析失败按未过期处理）
# ═══════════════════════════════════════════

def _to_aware(dt: datetime | None) -> datetime | None:
    """朴素时间视为 UTC，统一为 aware datetime（供跨时区安全比较）。"""
    if dt is None:
        return None
    if dt.tzinfo is None:
        return dt.replace(tzinfo=timezone.utc)
    return dt


def _parse_iso(value) -> datetime | None:
    """把 ISO8601 字符串（或 datetime）解析为 aware datetime，失败返回 None（不抛）。"""
    if value is None or value == "":
        return None
    if isinstance(value, datetime):
        return _to_aware(value)
    if not isinstance(value, str):
        return None
    try:
        s = value.strip()
        # 兼容末尾 Z（fromisoformat 3.11 前不识别）
        if s.endswith("Z"):
            s = s[:-1] + "+00:00"
        return _to_aware(datetime.fromisoformat(s))
    except Exception:
        return None


def _to_iso(value) -> str | None:
    """把 datetime/字符串归一化为 ISO 字符串（供前端展示），无值返回 None。"""
    if value is None:
        return None
    if isinstance(value, datetime):
        return _to_aware(value).isoformat()
    if isinstance(value, str) and value.strip():
        return value
    return None


def _compute_freshness(pd: dict, updated_at) -> dict:
    """计算单底稿新鲜度三态（P3 / Req1.2-1.4）。

    - checked_at = audit_checks_at 或 legacy fine_extracted_at（可空）
    - never_checked = checked_at 为空
    - stale = updated_at > checked_at（checked_at 为空或任一无法解析 → False，fail-safe）
    """
    checked_at_raw = pd.get("audit_checks_at") or pd.get("fine_extracted_at") or None
    never_checked = not bool(checked_at_raw)

    checked_dt = _parse_iso(checked_at_raw)
    updated_dt = _parse_iso(updated_at)

    if never_checked or checked_dt is None or updated_dt is None:
        stale = False
    else:
        stale = updated_dt > checked_dt

    return {
        "checked_at": _to_iso(checked_at_raw),
        "updated_at": _to_iso(updated_at),
        "stale": stale,
        "never_checked": never_checked,
    }


# ═══════════════════════════════════════════
# 检查项还原（缓存 dict → AuditCheckItem，供汇总）
# ═══════════════════════════════════════════

def _item_from_cache_dict(d: dict) -> AuditCheckItem:
    """把缓存 `audit_checks` 中的项 dict（`AuditCheckItem.to_dict()` 输出）还原为对象。

    兼容 legacy 键 `type` → `check_type`；源自 aggregator 同款还原逻辑（不改 aggregator，
    本端点自持一份轻量还原，仅用于汇总计算，字段读取容错）。
    """
    return AuditCheckItem(
        code=str(d.get("code", "")),
        source=str(d.get("source", "")),
        wp_code=str(d.get("wp_code", "")),
        wp_id=d.get("wp_id"),
        sheet_hint=d.get("sheet_hint"),
        severity=str(d.get("severity", SEVERITY_INFO)),
        check_type=str(d.get("check_type") or d.get("type", "")),
        description=str(d.get("description", "")),
        message=str(d.get("message", "")),
        produced_at=str(d.get("produced_at", "")),
        passed=d.get("passed"),
        actual=d.get("actual"),
        expected=d.get("expected"),
        diff=d.get("diff"),
    )


def _resolve_wp_checks(
    pd: dict, *, wp_code: str, wp_id: str
) -> tuple[list[dict], list[AuditCheckItem]]:
    """解析单底稿检查项，返回 (前端展示用 dict 列表, 汇总用 AuditCheckItem 列表)。

    向后兼容（P13）：
    - `parsed_data` 含 `audit_checks`（新字段，to_dict 超集结构）→ 直接用。
    - 否则退回 legacy `parsed_data.fine_checks`，用 `from_fine_check_dict` 补
      `source=fine_rule`/`wp_code`/`produced_at=fine_extracted_at` 转 dict（不报错）。
    - 均无 → 空（未检查底稿不贡献 check 项，P4）。
    """
    raw = pd.get("audit_checks")
    if isinstance(raw, list):
        check_dicts = [d for d in raw if isinstance(d, dict)]
        items = [_item_from_cache_dict(d) for d in check_dicts]
        return check_dicts, items

    # 退回 legacy fine_checks（补 source=fine_rule）
    fine_checks = pd.get("fine_checks", []) or []
    fine_extracted_at = str(pd.get("fine_extracted_at", "") or "")
    items = [
        from_fine_check_dict(
            fc, wp_code=wp_code, wp_id=wp_id, produced_at=fine_extracted_at
        )
        for fc in fine_checks
        if isinstance(fc, dict)
    ]
    check_dicts = [it.to_dict() for it in items]
    return check_dicts, items


# ═══════════════════════════════════════════
# 端点
# ═══════════════════════════════════════════

@router.get("/api/projects/{project_id}/audit-checks/summary")
async def get_audit_checks_summary(
    project_id: UUID,
    db: AsyncSession = Depends(get_db),
    current_user=Depends(require_project_access("readonly")),
):
    """批量获取项目所有底稿的统一审计检查结果 + 新鲜度 + 项目汇总。

    返回：
        {
          "summary": { total, decided, passed, failed, uncovered, pass_rate, blocking_open },
          "workpapers": [
            { wp_id, wp_code, wp_name, audit_cycle, checks:[...],
              checked_at, updated_at, stale, never_checked }
          ]
        }
    """
    # 一次批量 JOIN 查询（避免 N+1）
    wp_q = (
        sa.select(
            WorkingPaper.id,
            WorkingPaper.parsed_data,
            WorkingPaper.updated_at,
            WpIndex.wp_code,
            WpIndex.wp_name,
            WpIndex.audit_cycle,
        )
        .join(WpIndex, WorkingPaper.wp_index_id == WpIndex.id)
        .where(
            WorkingPaper.project_id == project_id,
            WorkingPaper.is_deleted == sa.false(),
        )
    )
    rows = (await db.execute(wp_q)).all()

    workpapers: list[dict] = []
    all_items: list[AuditCheckItem] = []

    for wp_id, parsed_data, updated_at, wp_code, wp_name, audit_cycle in rows:
        pd = parsed_data or {}
        wp_id_str = str(wp_id)

        check_dicts, items = _resolve_wp_checks(
            pd, wp_code=wp_code or "", wp_id=wp_id_str
        )
        all_items.extend(items)

        freshness = _compute_freshness(pd, updated_at)
        workpapers.append({
            "wp_id": wp_id_str,
            "wp_code": wp_code,
            "wp_name": wp_name,
            "audit_cycle": audit_cycle,
            "checks": check_dicts,
            **freshness,
        })

    summary = ProjectCheckSummary.from_items(all_items)

    return {
        "summary": summary.to_dict(),
        "workpapers": workpapers,
    }


# ═══════════════════════════════════════════
# 重算辅助
# ═══════════════════════════════════════════

async def _resolve_project_year(db: AsyncSession, project_id: UUID) -> int | None:
    """解析项目审计年度（优先 projects.audit_year，回退 audit_period_end 年份）。

    与 `cycle_review_context._get_project_year` 同口径（recompute_project 需要 year）。
    """
    row = (
        await db.execute(
            sa.select(
                Project.audit_year,
                sa.func.extract("year", Project.audit_period_end),
            ).where(Project.id == project_id)
        )
    ).first()
    if not row:
        return None
    audit_year, end_year = row
    if audit_year:
        return int(audit_year)
    return int(end_year) if end_year else None


async def _compute_project_summary(
    db: AsyncSession, project_id: UUID
) -> ProjectCheckSummary:
    """从持久化的 `parsed_data.audit_checks`（无则退回 legacy fine_checks）读全项目
    检查项，复用 summary 端点的解析逻辑计算项目汇总（重算后返回最新汇总用）。
    """
    rows = (
        await db.execute(
            sa.select(
                WorkingPaper.id,
                WorkingPaper.parsed_data,
                WpIndex.wp_code,
            )
            .join(WpIndex, WorkingPaper.wp_index_id == WpIndex.id)
            .where(
                WorkingPaper.project_id == project_id,
                WorkingPaper.is_deleted == sa.false(),
            )
        )
    ).all()

    all_items: list[AuditCheckItem] = []
    for wp_id, parsed_data, wp_code in rows:
        _dicts, items = _resolve_wp_checks(
            parsed_data or {}, wp_code=wp_code or "", wp_id=str(wp_id)
        )
        all_items.extend(items)

    return ProjectCheckSummary.from_items(all_items)


# ═══════════════════════════════════════════
# 重算端点（Req2.1-2.5 / 8.2）
# ═══════════════════════════════════════════

class RecomputeRequest(BaseModel):
    """重算请求体（可选 wp_id：传则只重算单张底稿，不传则重算项目全部）。"""

    wp_id: str | None = None


@router.post("/api/projects/{project_id}/audit-checks/recompute")
async def recompute_audit_checks(
    project_id: UUID,
    body: RecomputeRequest | None = Body(default=None),
    db: AsyncSession = Depends(get_db),
    current_user=Depends(require_project_access("edit")),
):
    """主动触发审计检查重算（Req2）。

    - 可选 body `{ "wp_id": "..." }`：传则只重算单张底稿（`recompute_workpaper`），
      不传则遍历项目全部底稿逐张重算。
    - 单底稿失败不阻断其余（Req2.4）：逐张 try/except 收集失败明细，继续其余；
      返回 `{recomputed, failed:[{wp_id,wp_code,reason}], summary}`。
    - 并发去重/幂等（Req2.3）：同项目重算进行中重复触发返回 `status="in_progress"`；
      重算为覆盖写天然幂等（重复触发结果一致、不累积脏数据）。
    - 权限（Req2.5）：编制权（`require_project_access("edit")`），无权限由依赖层拒绝。
    - 持久化：`recompute_workpaper` 内部已 flush 写 `parsed_data.audit_checks`；
      本端点负责 `await db.commit()`（参照 fine-extract 的 commit 模式）。

    Returns:
        {
          "status": "completed" | "in_progress",
          "recomputed": int,               # 成功重算的底稿数
          "failed": [{wp_id, wp_code, reason}],
          "summary": ProjectCheckSummary | None,  # 重算后的项目汇总
          "idempotent": true               # 重算覆盖写，重复触发结果一致
        }
    """
    pid_str = str(project_id)

    # 并发去重：进行中重复触发不重复启动（检查 + add 之间无 await → 原子）
    if pid_str in _recompute_in_progress:
        return {
            "status": "in_progress",
            "recomputed": 0,
            "failed": [],
            "summary": None,
            "idempotent": True,
            "message": "该项目审计检查重算正在进行中，请稍后刷新",
        }
    _recompute_in_progress.add(pid_str)

    try:
        year = await _resolve_project_year(db, project_id)
        aggregator = AuditCheckAggregator()
        failed: list[dict] = []
        recomputed = 0

        wp_id = body.wp_id if body else None

        if wp_id:
            # 单张底稿重算（需先查 wp + idx）
            try:
                wp_uuid = UUID(wp_id)
            except (ValueError, TypeError):
                raise HTTPException(status_code=400, detail="wp_id 格式非法")

            row = (
                await db.execute(
                    sa.select(WorkingPaper, WpIndex)
                    .join(WpIndex, WorkingPaper.wp_index_id == WpIndex.id)
                    .where(
                        WorkingPaper.id == wp_uuid,
                        WorkingPaper.project_id == project_id,
                        WorkingPaper.is_deleted == sa.false(),
                    )
                )
            ).first()
            if not row:
                raise HTTPException(status_code=404, detail="底稿不存在")
            wp, idx = row
            try:
                await aggregator.recompute_workpaper(db, wp, idx, year=year)
                recomputed = 1
            except Exception as exc:  # noqa: BLE001 — 单张失败如实记录不抛
                logger.warning(
                    "audit_check recompute single wp failed wp_id=%s", wp_id,
                    exc_info=True,
                )
                failed.append({
                    "wp_id": wp_id,
                    "wp_code": getattr(idx, "wp_code", ""),
                    "reason": str(exc) or exc.__class__.__name__,
                })
        else:
            # 项目全部底稿逐张重算（单张失败不阻断其余，Req2.4）
            rows = (
                await db.execute(
                    sa.select(WorkingPaper, WpIndex)
                    .join(WpIndex, WorkingPaper.wp_index_id == WpIndex.id)
                    .where(
                        WorkingPaper.project_id == project_id,
                        WorkingPaper.is_deleted == sa.false(),
                    )
                )
            ).all()
            for wp, idx in rows:
                try:
                    await aggregator.recompute_workpaper(db, wp, idx, year=year)
                    recomputed += 1
                except Exception as exc:  # noqa: BLE001 — 单张失败继续其余
                    logger.warning(
                        "audit_check recompute wp failed wp_id=%s",
                        getattr(wp, "id", "?"), exc_info=True,
                    )
                    failed.append({
                        "wp_id": str(getattr(wp, "id", "")),
                        "wp_code": getattr(idx, "wp_code", ""),
                        "reason": str(exc) or exc.__class__.__name__,
                    })

        # 持久化（参照 fine-extract 的 commit 模式；aggregator 内部已 flush）
        await db.commit()

        # 重算后的项目汇总（重新读持久化结果计算）
        summary = await _compute_project_summary(db, project_id)

        return {
            "status": "completed",
            "recomputed": recomputed,
            "failed": failed,
            "summary": summary.to_dict(),
            "idempotent": True,
        }
    finally:
        _recompute_in_progress.discard(pid_str)


# ═══════════════════════════════════════════
# 前端上报端点（Req3.1/3.3/4.1 — Task 4.2）
# ═══════════════════════════════════════════

class ReportedCheckItem(BaseModel):
    """前端上报的单条运行时勾稽检查项。

    由底稿现有 composable（tbReconcile/adjustmentReconcile/reportCrossCheck/crossSheet）
    的判定结果映射而来。`source`/`wp_code`/`wp_id`/`produced_at` 由服务端统一补齐
    （不信任前端传入），前端仅提供判定语义字段。
    """

    code: str
    severity: str = SEVERITY_INFO
    check_type: str = ""
    description: str = ""
    message: str = ""
    passed: bool | None = None  # 三态：True/False/None(未覆盖)
    actual: float | None = None
    expected: float | None = None
    diff: float | None = None
    sheet_hint: str | None = None


class ReportAuditChecksRequest(BaseModel):
    """上报请求体：一次上报一个 source 命名空间下的多条 item。

    顶层单 source + items（命名空间清晰）：同一 source 重报覆盖不累积（Property 6）。
    """

    source: str
    items: list[ReportedCheckItem] = []


@router.post(
    "/api/projects/{project_id}/workpapers/{wp_id}/audit-checks/report"
)
async def report_audit_checks(
    project_id: UUID,
    wp_id: UUID,
    body: ReportAuditChecksRequest = Body(...),
    db: AsyncSession = Depends(get_db),
    current_user=Depends(require_project_access("edit")),
):
    """前端上报运行时勾稽真源（S6），按 `source` 命名空间 upsert 到该底稿的
    `parsed_data.audit_checks`（与后端自算 S1-S5 写同一缓存，summary 一处读）。

    行为（design §3 + Property 6 / Req4.1 / Req3.3）：
    - **校验 source（契约）**：`source` 必须 ∈ `FRONTEND_REPORTABLE_SOURCES`
      （tb_recon/adjustment_recon/report_cross_check/cross_sheet），否则 400；
      特别地 `source ∈ BACKEND_COMPUTED_SOURCES`（前端伪造后端专属 source）→ 400 拒绝。
      校验/归属校验置于 try 外，HTTPException 不被通用 except 吞成 500。
    - **wp 归属校验**：`WorkingPaper JOIN WpIndex WHERE id==wp_id AND
      project_id==project_id AND is_deleted==false`，不存在 → 404。
    - **upsert 语义**：移除该 wp 上所有 `source == 本次上报 source` 的旧项（同 source 覆盖
      不累积），再把本次 items 追加；**不动其他 source 的项**（含后端自算 S1-S5 与其他
      S6 source）。每条补齐 `source`/`wp_code`（从 idx 查）/`wp_id`/`produced_at=now`；
      `passed` 保留前端三态。
    - 同步更新 `audit_checks_at=now`（让 summary 端点的 stale 新鲜度正确反映最近变更）。

    权限（Req4.1）：编制权（`require_project_access("edit")`，注入即校验）。

    Returns:
        { "reported": int, "source": str, "total_checks": int }
    """
    # ── 校验 source（置于 try 外，HTTPException 不被吞成 500）──────────────
    source = body.source
    if source in BACKEND_COMPUTED_SOURCES or source not in FRONTEND_REPORTABLE_SOURCES:
        raise HTTPException(status_code=400, detail=f"不支持的上报来源: {source}")

    # ── wp 归属校验（不存在 → 404，置于 try 外）──────────────────────────
    row = (
        await db.execute(
            sa.select(WorkingPaper, WpIndex)
            .join(WpIndex, WorkingPaper.wp_index_id == WpIndex.id)
            .where(
                WorkingPaper.id == wp_id,
                WorkingPaper.project_id == project_id,
                WorkingPaper.is_deleted == sa.false(),
            )
        )
    ).first()
    if not row:
        raise HTTPException(status_code=404, detail="底稿不存在")
    wp, idx = row

    wp_code = getattr(idx, "wp_code", "") or ""
    wp_id_str = str(wp.id)
    now_iso = datetime.now(timezone.utc).isoformat()

    # ── 本次上报项转 AuditCheckItem.to_dict()（服务端统一补齐归属/时间）───────
    new_dicts: list[dict] = []
    for it in body.items:
        item = AuditCheckItem(
            code=str(it.code or ""),
            source=source,
            wp_code=wp_code,
            wp_id=wp_id_str,
            sheet_hint=it.sheet_hint,
            severity=str(it.severity or SEVERITY_INFO),
            check_type=str(it.check_type or ""),
            description=str(it.description or ""),
            message=str(it.message or ""),
            produced_at=now_iso,
            passed=it.passed,  # 保留前端三态 True/False/None
            actual=it.actual,
            expected=it.expected,
            diff=it.diff,
        )
        new_dicts.append(item.to_dict())

    try:
        pd = wp.parsed_data or {}
        existing = pd.get("audit_checks")
        existing = [d for d in existing if isinstance(d, dict)] if isinstance(existing, list) else []

        # upsert：移除本次 source 的旧项（同 source 覆盖不累积），保留其他 source（S1-S5/其他 S6）
        kept = [d for d in existing if d.get("source") != source]
        merged = kept + new_dicts

        pd["audit_checks"] = merged
        pd["audit_checks_at"] = now_iso  # 上报也刷新变更时间，让 stale 正确反映
        wp.parsed_data = pd
        flag_modified(wp, "parsed_data")
        await db.flush()
        await db.commit()
    except HTTPException:
        raise
    except Exception:  # noqa: BLE001
        await db.rollback()
        logger.warning(
            "audit_check report upsert failed wp_id=%s source=%s", wp_id_str, source,
            exc_info=True,
        )
        raise HTTPException(status_code=500, detail="上报审计检查失败")

    return {
        "reported": len(new_dicts),
        "source": source,
        "total_checks": len(merged),
    }


# ═══════════════════════════════════════════
# 签认端点（Req8 — Task 5.2，只提示不阻断 P14 + 快照一致 P10）
# ═══════════════════════════════════════════

class SignoffRequest(BaseModel):
    """复核签认请求体（note 可选备注）。"""

    note: str | None = None


def _resolve_signoff_year(year: int | None) -> int:
    """签认年度：优先项目审计年度，缺失时回退当前自然年（year 列 NOT NULL）。"""
    if year:
        return int(year)
    return datetime.now(timezone.utc).year


@router.post("/api/projects/{project_id}/audit-checks/signoff")
async def create_audit_check_signoff(
    project_id: UUID,
    body: SignoffRequest | None = Body(default=None),
    db: AsyncSession = Depends(get_db),
    current_user=Depends(require_project_access("review")),
):
    """复核签认——把项目概览「审计检查」当作复核最后一次检查的签认留痕（Req8）。

    行为（design §3 / Req8.3 / P10 / P14）：
    1. 算当前项目汇总 `summary`（复用 `_compute_project_summary`）。
    2. `blocking_present = summary.blocking_open > 0`。
    3. **只提示不阻断（P14）**：即使 `blocking_present=True` 也**照常写签认记录并返回成功**，
       不 raise / 不拒绝——仅在响应 message 中提示存在未处理阻断项。
    4. 写 `AuditCheckSignoff`：project_id / year（`_resolve_project_year`）/ signed_by
       （`current_user.id`）/ signed_by_name（`current_user.username`）/
       `summary_snapshot = summary.to_dict()`（P10 快照与当时 summary 一致）/
       blocking_present / note（body 可选）。commit。

    权限（Req8.4）：复核权（`require_project_access("review")`，高于只读，注入即校验）。

    Returns:
        {
          signoff_id, signed_at, blocking_present,
          summary: ProjectCheckSummary,
          message: str
        }
    """
    summary = await _compute_project_summary(db, project_id)
    blocking_present = summary.blocking_open > 0
    year = _resolve_signoff_year(await _resolve_project_year(db, project_id))

    signoff = AuditCheckSignoff(
        project_id=project_id,
        year=year,
        signed_by=current_user.id,
        signed_by_name=getattr(current_user, "username", None),
        summary_snapshot=summary.to_dict(),  # 快照与当时 summary 完全一致（P10）
        blocking_present=blocking_present,
        note=(body.note if body and body.note else None),
    )
    db.add(signoff)
    await db.commit()
    await db.refresh(signoff)  # 取 server_default 生成的 id / signed_at

    return {
        "signoff_id": str(signoff.id),
        "signed_at": _to_iso(signoff.signed_at),
        "blocking_present": blocking_present,
        "summary": summary.to_dict(),
        "message": (
            "存在未处理阻断项，已记录签认（仅提示）"
            if blocking_present
            else "签认成功"
        ),
    }


@router.get("/api/projects/{project_id}/audit-checks/signoff")
async def get_latest_audit_check_signoff(
    project_id: UUID,
    db: AsyncSession = Depends(get_db),
    current_user=Depends(require_project_access("readonly")),
):
    """读最近一次复核签认（Req8）。

    `WHERE project_id AND is_deleted=false ORDER BY signed_at DESC LIMIT 1`。
    无则返回 `{ "signoff": null }`；有则返回该记录字段。

    权限：项目只读（`require_project_access("readonly")`）。
    """
    row = (
        await db.execute(
            sa.select(AuditCheckSignoff)
            .where(
                AuditCheckSignoff.project_id == project_id,
                AuditCheckSignoff.is_deleted == sa.false(),
            )
            .order_by(AuditCheckSignoff.signed_at.desc())
            .limit(1)
        )
    ).scalar_one_or_none()

    if row is None:
        return {"signoff": None}

    return {
        "signoff": {
            "id": str(row.id),
            "signed_by": str(row.signed_by),
            "signed_by_name": row.signed_by_name,
            "signed_at": _to_iso(row.signed_at),
            "summary_snapshot": row.summary_snapshot,
            "blocking_present": row.blocking_present,
            "note": row.note,
        }
    }


# ═══════════════════════════════════════════
# 导出端点（Req7 — Task 5.3，xlsx + RFC5987 中文名 + 空数据不报错）
# ═══════════════════════════════════════════

# source → 中文标签（导出明细展示；与前端 SOURCE_META 同口径）
_SOURCE_LABELS: dict[str, str] = {
    "fine_rule": "精细化规则",
    "cycle_recon": "审定勾稽",
    "note_validation": "附注校验",
    "qc": "质控",
    "unadjusted_misstatement": "未更正错报",
    "tb_recon": "审定↔TB",
    "adjustment_recon": "审定↔调整",
    "report_cross_check": "报表核对",
    "cross_sheet": "跨表",
}

_SEVERITY_LABELS_EXPORT: dict[str, str] = {
    "blocking": "阻断",
    "warning": "警告",
    "info": "提示",
}


def _passed_label(passed) -> str:
    """判定三态 → 中文（True=通过 / False=未通过 / None=未覆盖）。"""
    if passed is True:
        return "通过"
    if passed is False:
        return "未通过"
    return "未覆盖"


@router.post("/api/projects/{project_id}/audit-checks/export")
async def export_audit_checks(
    project_id: UUID,
    db: AsyncSession = Depends(get_db),
    current_user=Depends(require_project_access("readonly")),
):
    """导出审计检查结果为 xlsx（Req7，留痕）。

    - **Sheet1「检查明细」**：底稿编码/底稿名称/检查编号/来源/严重程度/类型/判定/消息，
      逐底稿逐 check 一行。
    - **Sheet2「汇总」**：项目汇总（total/decided/passed/failed/uncovered/pass_rate/
      blocking_open）+ 导出时间。
    - **无检查数据不报错**（Req7 Error Handling）：导出空模板（表头齐全、无数据行）+
      汇总 decided=0，正常返回。
    - **RFC5987 中文文件名**：`审计检查结果_{项目}_{年度}.xlsx`（平台铁律，
      复用 `content_disposition_attachment`）。

    权限：项目只读（`require_project_access("readonly")`，导出为只读取数动作）。
    """
    from io import BytesIO

    from fastapi.responses import StreamingResponse

    try:
        import openpyxl
        from openpyxl.styles import Alignment, Font, PatternFill
    except ImportError:
        raise HTTPException(status_code=500, detail="openpyxl 未安装")

    # 项目名 + 年度（供文件名）
    proj_row = (
        await db.execute(
            sa.select(Project.name, Project.client_name).where(Project.id == project_id)
        )
    ).first()
    project_label = ""
    if proj_row:
        project_label = (proj_row[1] or proj_row[0] or "").strip()
    year = _resolve_signoff_year(await _resolve_project_year(db, project_id))

    # 批量读全项目底稿检查项（复用 summary 端点的解析逻辑）
    rows = (
        await db.execute(
            sa.select(
                WorkingPaper.id,
                WorkingPaper.parsed_data,
                WpIndex.wp_code,
                WpIndex.wp_name,
            )
            .join(WpIndex, WorkingPaper.wp_index_id == WpIndex.id)
            .where(
                WorkingPaper.project_id == project_id,
                WorkingPaper.is_deleted == sa.false(),
            )
        )
    ).all()

    all_items: list[AuditCheckItem] = []
    detail_rows: list[tuple] = []  # (wp_code, wp_name, code, source, severity, type, 判定, message)
    for wp_id, parsed_data, wp_code, wp_name in rows:
        _dicts, items = _resolve_wp_checks(
            parsed_data or {}, wp_code=wp_code or "", wp_id=str(wp_id)
        )
        all_items.extend(items)
        for it in items:
            detail_rows.append((
                wp_code or "",
                wp_name or "",
                it.code,
                _SOURCE_LABELS.get(it.source, it.source or "其他"),
                _SEVERITY_LABELS_EXPORT.get(it.severity, it.severity),
                it.check_type,
                _passed_label(it.passed),
                it.message,
            ))

    summary = ProjectCheckSummary.from_items(all_items)

    # ── 构建工作簿 ──
    header_font = Font(bold=True, color="FFFFFF")
    header_fill = PatternFill(start_color="4B2D77", end_color="4B2D77", fill_type="solid")
    center = Alignment(horizontal="center", vertical="center")

    wb = openpyxl.Workbook()

    # Sheet1 检查明细（无数据仍写表头 → 空模板）
    ws1 = wb.active
    ws1.title = "检查明细"
    detail_headers = ["底稿编码", "底稿名称", "检查编号", "来源", "严重程度", "类型", "判定", "消息"]
    ws1.append(detail_headers)
    for ci in range(1, len(detail_headers) + 1):
        c = ws1.cell(row=1, column=ci)
        c.font = header_font
        c.fill = header_fill
        c.alignment = center
    for r in detail_rows:
        ws1.append(list(r))
    for idx_col, width in enumerate([14, 24, 16, 14, 12, 14, 10, 48], start=1):
        ws1.column_dimensions[openpyxl.utils.get_column_letter(idx_col)].width = width

    # Sheet2 汇总
    ws2 = wb.create_sheet("汇总")
    pass_rate_display = (
        f"{round(summary.pass_rate * 100, 2)}%" if summary.pass_rate is not None else "—"
    )
    summary_rows = [
        ("总检查数", summary.total),
        ("已判定数", summary.decided),
        ("通过", summary.passed),
        ("未通过", summary.failed),
        ("未覆盖", summary.uncovered),
        ("已判定通过率", pass_rate_display),
        ("未处理阻断项", summary.blocking_open),
        ("导出时间", datetime.now(timezone.utc).strftime("%Y-%m-%d %H:%M:%S UTC")),
    ]
    ws2.append(["指标", "值"])
    for ci in (1, 2):
        c = ws2.cell(row=1, column=ci)
        c.font = header_font
        c.fill = header_fill
        c.alignment = center
    for label, value in summary_rows:
        ws2.append([label, value])
    ws2.column_dimensions["A"].width = 18
    ws2.column_dimensions["B"].width = 28

    buf = BytesIO()
    wb.save(buf)
    buf.seek(0)

    fname_cn = f"审计检查结果_{project_label}_{year}.xlsx" if project_label else f"审计检查结果_{year}.xlsx"
    return StreamingResponse(
        buf,
        media_type="application/vnd.openxmlformats-officedocument.spreadsheetml.sheet",
        headers={
            "Content-Disposition": content_disposition_attachment(
                fname_cn,
                ascii_fallback=f"audit_checks_{year}.xlsx",
            ),
        },
    )

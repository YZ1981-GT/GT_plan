"""N3 递延所得税负债 — 导入导出三级端点

3个端点：
- GET  /api/n3-deferred-tax-liabilities/{wp_id}/export-template   空白模板xlsx
- GET  /api/n3-deferred-tax-liabilities/{wp_id}/export-data        当前数据xlsx
- POST /api/n3-deferred-tax-liabilities/{wp_id}/import-data        解析xlsx写入（multipart）

科目编码: 2901递延所得税负债（**贷方/负债类**）
导入导出目标: N3-2明细表（动态行）
列: 项目名称 / 分类 / 账面价值 / 计税基础 / 适用税率 / 期初递延税负债 / 本期确认 / 本期转回 / 备注
RFC5987 Content-Disposition header with Chinese filename encoding

Requirements: 3.4
"""

from __future__ import annotations

import io
import json
import logging
from typing import Any
from urllib.parse import quote
from uuid import uuid4

from fastapi import APIRouter, Depends, File, HTTPException, Query, UploadFile
from fastapi.responses import StreamingResponse
from openpyxl import Workbook, load_workbook
from openpyxl.styles import Alignment, Font, PatternFill
from sqlalchemy.ext.asyncio import AsyncSession
import sqlalchemy as sa

from app.core.database import get_db
from app.deps import get_current_user
from app.models.core import User

logger = logging.getLogger(__name__)

router = APIRouter(
    prefix="/api/n3-deferred-tax-liabilities",
    tags=["N3 递延所得税负债"],
)

# ═══════════════════════════════════════════════════════════════════════════════
# Sheet 配置
# ═══════════════════════════════════════════════════════════════════════════════

_ROW_LIMIT = 500

_SUPPORTED_SHEETS: set[str] = {"N3-2"}

# N3-2 明细表列头（9列，对应应纳税暂时性差异明细）
_SHEET_HEADERS: dict[str, list[str]] = {
    "N3-2": [
        "项目名称", "分类", "账面价值", "计税基础",
        "适用税率", "期初递延税负债", "本期确认", "本期转回", "备注",
    ],
}

# 中文列名 → JSON 字段名映射
_FIELD_MAPS: dict[str, dict[str, str]] = {
    "N3-2": {
        "项目名称": "projectName",
        "分类": "category",
        "账面价值": "bookValue",
        "计税基础": "taxBase",
        "适用税率": "taxRate",
        "期初递延税负债": "openingDtl",
        "本期确认": "currentRecognized",
        "本期转回": "currentReversed",
        "备注": "remark",
    },
}

# 数值类型字段
_NUMERIC_FIELDS: dict[str, set[str]] = {
    "N3-2": {
        "bookValue", "taxBase", "taxRate",
        "openingDtl", "currentRecognized", "currentReversed",
    },
}

# checklist_responses item_id
_SHEET_ITEM_ID: dict[str, str] = {
    "N3-2": "N3-detail-rows",
}


# ═══════════════════════════════════════════════════════════════════════════════
# 辅助函数
# ═══════════════════════════════════════════════════════════════════════════════


def _validate_sheet(sheet_code: str) -> None:
    """校验sheet编码"""
    if sheet_code not in _SUPPORTED_SHEETS:
        raise HTTPException(
            400,
            f"不支持的sheet: {sheet_code}。支持: {sorted(_SUPPORTED_SHEETS)}",
        )


def _safe_float(val: Any) -> float:
    """安全转float"""
    if val is None:
        return 0.0
    try:
        return float(val)
    except (ValueError, TypeError):
        return 0.0


def _safe_str(val: Any) -> str:
    """安全转str"""
    if val is None:
        return ""
    return str(val).strip()


def _create_template_wb(sheet_code: str) -> Workbook:
    """创建空白模板xlsx（含表头+格式，无数据行）"""
    wb = Workbook()
    ws = wb.active
    ws.title = sheet_code

    headers = _SHEET_HEADERS[sheet_code]
    header_font = Font(bold=True, size=11)
    header_fill = PatternFill(start_color="D9E1F2", end_color="D9E1F2", fill_type="solid")
    header_align = Alignment(horizontal="center", vertical="center", wrap_text=True)

    for col_idx, col_name in enumerate(headers, 1):
        cell = ws.cell(row=1, column=col_idx, value=col_name)
        cell.font = header_font
        cell.fill = header_fill
        cell.alignment = header_align
        ws.column_dimensions[cell.column_letter].width = max(len(col_name) * 2 + 4, 12)

    ws.freeze_panes = "A2"
    return wb


def _get_actual_headers(ws: Any) -> list[str]:
    """获取worksheet实际列头"""
    headers: list[str] = []
    for cell in next(ws.iter_rows(min_row=1, max_row=1)):
        if cell.value is not None:
            headers.append(str(cell.value).strip())
    return headers


def _validate_columns(ws: Any, sheet_code: str) -> list[str]:
    """校验列头，返回不匹配的列名列表"""
    expected = set(_SHEET_HEADERS[sheet_code])
    actual = set(_get_actual_headers(ws))
    missing = [h for h in expected if h not in actual]
    return missing


def _col_val(row: tuple, actual_headers: list[str], col_name: str) -> Any:
    """从行数据中按列名取值"""
    try:
        idx = actual_headers.index(col_name)
        return row[idx] if idx < len(row) else None
    except (ValueError, IndexError):
        return None


def _export_row(sheet_code: str, data: dict) -> list:
    """将JSON dict按列头顺序转为导出行"""
    headers = _SHEET_HEADERS[sheet_code]
    field_map = _FIELD_MAPS.get(sheet_code, {})
    result = []
    for col_name in headers:
        field = field_map.get(col_name, col_name)
        val = data.get(field, "")
        result.append(val if val is not None else "")
    return result


def _parse_row(sheet_code: str, row: tuple, actual_headers: list[str]) -> dict:
    """将xlsx行数据按列名映射为JSON dict"""
    field_map = _FIELD_MAPS.get(sheet_code, {})
    result: dict[str, Any] = {"rowId": str(uuid4())}
    for col_name, field_name in field_map.items():
        raw = _col_val(row, actual_headers, col_name)
        if field_name in _NUMERIC_FIELDS.get(sheet_code, set()):
            result[field_name] = _safe_float(raw)
        else:
            result[field_name] = _safe_str(raw)
    return result


# ═══════════════════════════════════════════════════════════════════════════════
# 端点
# ═══════════════════════════════════════════════════════════════════════════════


@router.get("/{wp_id}/export-template")
async def n3_export_template(
    wp_id: str,
    sheet: str = Query("N3-2", description="Sheet编码: N3-2"),
    current_user: User = Depends(get_current_user),
) -> StreamingResponse:
    """导出空白xlsx模板（含表头+格式，无数据行）

    Query params:
        sheet: 指定导出sheet（默认N3-2）

    Requirements: 3.4
    """
    _validate_sheet(sheet)

    wb = _create_template_wb(sheet)
    buffer = io.BytesIO()
    wb.save(buffer)
    buffer.seek(0)

    filename = f"N3递延所得税负债_{sheet}_模板.xlsx"
    return StreamingResponse(
        buffer,
        media_type="application/vnd.openxmlformats-officedocument.spreadsheetml.sheet",
        headers={"Content-Disposition": f"attachment; filename*=UTF-8''{quote(filename)}"},
    )


@router.get("/{wp_id}/export-data")
async def n3_export_data(
    wp_id: str,
    sheet: str = Query("N3-2", description="Sheet编码: N3-2"),
    db: AsyncSession = Depends(get_db),
    current_user: User = Depends(get_current_user),
) -> StreamingResponse:
    """导出含当前数据的xlsx（N3-2明细表动态行数据）

    Query params:
        sheet: 指定导出sheet（默认N3-2）

    Requirements: 3.4
    """
    _validate_sheet(sheet)

    item_id = _SHEET_ITEM_ID[sheet]
    rows_data: list[dict] = []

    result = await db.execute(
        sa.text(
            "SELECT remark FROM checklist_responses "
            "WHERE wp_id = :wp_id AND item_id = :item_id LIMIT 1"
        ),
        {"wp_id": wp_id, "item_id": item_id},
    )
    row = result.fetchone()
    if row and row.remark:
        try:
            rows_data = json.loads(row.remark)
        except (json.JSONDecodeError, TypeError):
            pass

    # 生成 xlsx
    wb = _create_template_wb(sheet)
    ws = wb.active

    for data_row in rows_data:
        ws.append(_export_row(sheet, data_row))

    buffer = io.BytesIO()
    wb.save(buffer)
    buffer.seek(0)

    filename = f"N3递延所得税负债_{sheet}_数据.xlsx"
    return StreamingResponse(
        buffer,
        media_type="application/vnd.openxmlformats-officedocument.spreadsheetml.sheet",
        headers={"Content-Disposition": f"attachment; filename*=UTF-8''{quote(filename)}"},
    )


@router.post("/{wp_id}/import-data")
async def n3_import_data(
    wp_id: str,
    sheet: str = Query("N3-2", description="Sheet编码: N3-2"),
    file: UploadFile = File(...),
    db: AsyncSession = Depends(get_db),
    current_user: User = Depends(get_current_user),
) -> dict[str, Any]:
    """导入xlsx解析写入checklist_responses（N3-2明细表动态行）

    Query params:
        sheet: 指定导入目标sheet（默认N3-2）

    验证：文件大小 ≤ 10MB、文件扩展名(.xlsx/.xls)、列头匹配

    Returns:
        { imported_count: int, field_count: int, warning?: str }

    Requirements: 3.4
    """
    _validate_sheet(sheet)

    # 读取上传文件
    content = await file.read()
    if len(content) > 10 * 1024 * 1024:
        raise HTTPException(400, "文件大小超过 10MB 限制")

    # 文件扩展名校验
    filename = file.filename or ""
    suffix = filename.rsplit(".", 1)[-1].lower() if "." in filename else ""
    if suffix not in ("xlsx", "xls"):
        raise HTTPException(400, f"不支持的文件类型: .{suffix}，仅支持 .xlsx/.xls")

    try:
        wb = load_workbook(io.BytesIO(content), read_only=True, data_only=True)
    except Exception:
        raise HTTPException(400, "无法解析xlsx文件")

    ws = wb.active

    # 校验列头
    missing = _validate_columns(ws, sheet)
    if missing:
        raise HTTPException(400, f"列名不匹配，缺少: {missing}")

    actual_headers = _get_actual_headers(ws)

    # 解析数据行
    parsed_rows: list[dict] = []
    warning: str | None = None
    for row_idx, row in enumerate(ws.iter_rows(min_row=2, values_only=True), start=2):
        if row_idx > _ROW_LIMIT + 1:
            warning = f"数据行超过{_ROW_LIMIT}行限制，已截断"
            break
        # 跳过全空行
        if all(cell is None or str(cell).strip() == "" for cell in row):
            continue
        parsed_rows.append(_parse_row(sheet, row, actual_headers))

    # 写入 checklist_responses
    item_id = _SHEET_ITEM_ID[sheet]
    json_str = json.dumps(parsed_rows, ensure_ascii=False)

    await db.execute(
        sa.text(
            "INSERT INTO checklist_responses (id, wp_id, item_id, remark) "
            "VALUES (:id, :wp_id, :item_id, :remark) "
            "ON CONFLICT (wp_id, item_id) DO UPDATE SET remark = :remark"
        ),
        {
            "id": str(uuid4()),
            "wp_id": wp_id,
            "item_id": item_id,
            "remark": json_str,
        },
    )
    await db.commit()

    field_count = len(_FIELD_MAPS.get(sheet, {}))
    result: dict[str, Any] = {
        "imported_count": len(parsed_rows),
        "field_count": field_count,
        "message": f"成功导入 {len(parsed_rows)} 行数据",
    }
    if warning:
        result["warning"] = warning
    return result

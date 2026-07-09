"""N5 所得税费用 — 导入导出三级端点

3个端点：
- GET  /api/n5-income-tax-expense/{wp_id}/export-template   空白模板xlsx（纳税调整分sheet）
- GET  /api/n5-income-tax-expense/{wp_id}/export-data        当前数据xlsx（纳税调整分sheet）
- POST /api/n5-income-tax-expense/{wp_id}/import-data        解析xlsx写入（multipart）

科目编码: 6801所得税费用（**损益类**！取本期发生额）
导入导出目标: N5-2明细表（动态行）/ N5-5纳税调整明细（107行动态行，分sheet导出）
RFC5987 Content-Disposition header with Chinese filename encoding

Requirements: 4.6
"""

from __future__ import annotations

import io
import json
import logging
from pathlib import Path
from typing import Any
from urllib.parse import quote
from uuid import uuid4

from fastapi import APIRouter, Depends, File, HTTPException, Query, UploadFile
from fastapi.responses import StreamingResponse
from openpyxl import Workbook, load_workbook
from openpyxl.styles import Alignment, Font, PatternFill
from sqlalchemy.ext.asyncio import AsyncSession
import sqlalchemy as sa

from app.core.database import get_db
from app.deps import get_current_user
from app.models.core import User

logger = logging.getLogger(__name__)

router = APIRouter(
    prefix="/api/n5-income-tax-expense",
    tags=["N5 所得税费用"],
)

# ═══════════════════════════════════════════════════════════════════════════════
# Sheet 配置
# ═══════════════════════════════════════════════════════════════════════════════

_ROW_LIMIT = 500

# 支持导入导出的sheet（动态行表格）
_SUPPORTED_SHEETS: set[str] = {"N5-2", "N5-5"}

# N5-2 明细表列头（10列，当期/递延分项明细）
# N5-5 纳税调整明细表列头（8列，107行调增/调减）
_SHEET_HEADERS: dict[str, list[str]] = {
    "N5-2": [
        "项目名称", "分类", "本期发生额", "未审数",
        "AJE调整", "RJE调整", "审定数", "上期数", "差异", "备注",
    ],
    "N5-5": [
        "序号", "纳税调整项目", "分类", "账载金额",
        "税收金额", "调增金额", "调减金额", "依据",
    ],
}

# 中文列名 → JSON 字段名映射
_FIELD_MAPS: dict[str, dict[str, str]] = {
    "N5-2": {
        "项目名称": "projectName",
        "分类": "category",
        "本期发生额": "periodAmount",
        "未审数": "unadjusted",
        "AJE调整": "aje",
        "RJE调整": "rje",
        "审定数": "audited",
        "上期数": "priorPeriod",
        "差异": "difference",
        "备注": "remark",
    },
    "N5-5": {
        "序号": "seqNo",
        "纳税调整项目": "adjustmentItem",
        "分类": "category",
        "账载金额": "bookAmount",
        "税收金额": "taxAmount",
        "调增金额": "addBackAmount",
        "调减金额": "deductAmount",
        "依据": "basis",
    },
}

# 数值类型字段
_NUMERIC_FIELDS: dict[str, set[str]] = {
    "N5-2": {
        "periodAmount", "unadjusted", "aje", "rje",
        "audited", "priorPeriod", "difference",
    },
    "N5-5": {
        "seqNo", "bookAmount", "taxAmount",
        "addBackAmount", "deductAmount",
    },
}

# checklist_responses item_id 前缀
_SHEET_ITEM_ID: dict[str, str] = {
    "N5-2": "N5-detail-rows",
    "N5-5": "N5-tax-adjustment-rows",
}

# N5-5 纳税调整分类（分sheet导出时使用）
_TAX_ADJUSTMENT_CATEGORIES = [
    "收入类调整项目",
    "扣除类调整项目",
    "资产类调整项目",
    "特殊事项调整项目",
    "其他",
]


# ═══════════════════════════════════════════════════════════════════════════════
# 辅助函数
# ═══════════════════════════════════════════════════════════════════════════════


def _validate_sheet(sheet_code: str) -> None:
    """校验sheet编码"""
    if sheet_code not in _SUPPORTED_SHEETS:
        raise HTTPException(
            400,
            f"不支持的sheet: {sheet_code}。支持: {sorted(_SUPPORTED_SHEETS)}",
        )


def _safe_float(val: Any) -> float:
    """安全转float"""
    if val is None:
        return 0.0
    try:
        return float(val)
    except (ValueError, TypeError):
        return 0.0


def _safe_str(val: Any) -> str:
    """安全转str"""
    if val is None:
        return ""
    return str(val).strip()


def _create_template_wb(sheet_code: str) -> Workbook:
    """创建空白模板xlsx（含表头+格式，无数据行）"""
    wb = Workbook()
    ws = wb.active
    ws.title = sheet_code

    headers = _SHEET_HEADERS[sheet_code]
    header_font = Font(bold=True, size=11)
    header_fill = PatternFill(start_color="D9E1F2", end_color="D9E1F2", fill_type="solid")
    header_align = Alignment(horizontal="center", vertical="center", wrap_text=True)

    for col_idx, col_name in enumerate(headers, 1):
        cell = ws.cell(row=1, column=col_idx, value=col_name)
        cell.font = header_font
        cell.fill = header_fill
        cell.alignment = header_align
        ws.column_dimensions[cell.column_letter].width = max(len(col_name) * 2 + 4, 12)

    ws.freeze_panes = "A2"
    return wb


def _create_n5_5_split_template() -> Workbook:
    """创建N5-5纳税调整明细分sheet模板（按分类拆分为多sheet）"""
    wb = Workbook()
    # 删除默认Sheet
    wb.remove(wb.active)

    headers = _SHEET_HEADERS["N5-5"]
    header_font = Font(bold=True, size=11)
    header_fill = PatternFill(start_color="D9E1F2", end_color="D9E1F2", fill_type="solid")
    header_align = Alignment(horizontal="center", vertical="center", wrap_text=True)

    for cat_name in _TAX_ADJUSTMENT_CATEGORIES:
        ws = wb.create_sheet(title=cat_name[:31])  # Excel sheet名最长31字符
        for col_idx, col_name in enumerate(headers, 1):
            cell = ws.cell(row=1, column=col_idx, value=col_name)
            cell.font = header_font
            cell.fill = header_fill
            cell.alignment = header_align
            ws.column_dimensions[cell.column_letter].width = max(len(col_name) * 2 + 4, 12)
        ws.freeze_panes = "A2"

    return wb


def _get_actual_headers(ws: Any) -> list[str]:
    """获取worksheet实际列头"""
    headers: list[str] = []
    for cell in next(ws.iter_rows(min_row=1, max_row=1)):
        if cell.value is not None:
            headers.append(str(cell.value).strip())
    return headers


def _validate_columns(ws: Any, sheet_code: str) -> list[str]:
    """校验列头，返回不匹配的列名列表"""
    expected = set(_SHEET_HEADERS[sheet_code])
    actual = set(_get_actual_headers(ws))
    missing = [h for h in expected if h not in actual]
    return missing


def _col_val(row: tuple, actual_headers: list[str], col_name: str) -> Any:
    """从行数据中按列名取值"""
    try:
        idx = actual_headers.index(col_name)
        return row[idx] if idx < len(row) else None
    except (ValueError, IndexError):
        return None


def _export_row(sheet_code: str, data: dict) -> list:
    """将JSON dict按列头顺序转为导出行"""
    headers = _SHEET_HEADERS[sheet_code]
    field_map = _FIELD_MAPS.get(sheet_code, {})
    result = []
    for col_name in headers:
        field = field_map.get(col_name, col_name)
        val = data.get(field, "")
        result.append(val if val is not None else "")
    return result


def _parse_row(sheet_code: str, row: tuple, actual_headers: list[str]) -> dict:
    """将xlsx行数据按列名映射为JSON dict"""
    field_map = _FIELD_MAPS.get(sheet_code, {})
    result: dict[str, Any] = {"rowId": str(uuid4())}
    for col_name, field_name in field_map.items():
        raw = _col_val(row, actual_headers, col_name)
        if field_name in _NUMERIC_FIELDS.get(sheet_code, set()):
            result[field_name] = _safe_float(raw)
        else:
            result[field_name] = _safe_str(raw)
    return result


async def _load_rows_data(wp_id: str, sheet_code: str, db: AsyncSession) -> list[dict]:
    """从checklist_responses加载动态行数据"""
    item_id = _SHEET_ITEM_ID[sheet_code]
    result = await db.execute(
        sa.text(
            "SELECT remark FROM checklist_responses "
            "WHERE wp_id = :wp_id AND item_id = :item_id LIMIT 1"
        ),
        {"wp_id": wp_id, "item_id": item_id},
    )
    row = result.fetchone()
    if row and row.remark:
        try:
            return json.loads(row.remark)
        except (json.JSONDecodeError, TypeError):
            pass
    return []


# ═══════════════════════════════════════════════════════════════════════════════
# 导出模板
# ═══════════════════════════════════════════════════════════════════════════════


@router.get("/{wp_id}/export-template")
async def n5_export_template(
    wp_id: str,
    sheet: str = Query("N5-5", description="Sheet编码: N5-2 / N5-5"),
    current_user: User = Depends(get_current_user),
) -> StreamingResponse:
    """导出空白xlsx模板（纳税调整分sheet）

    Query params:
        sheet: 指定导出sheet（默认N5-5）
            - N5-2: 所得税费用明细表
            - N5-5: 纳税调整明细表（分5分类sheet导出）

    Requirements: 4.6
    """
    _validate_sheet(sheet)

    if sheet == "N5-5":
        # N5-5 纳税调整：分类拆分为多sheet模板
        wb = _create_n5_5_split_template()
    else:
        wb = _create_template_wb(sheet)

    buffer = io.BytesIO()
    wb.save(buffer)
    buffer.seek(0)

    filename = f"N5所得税费用_{sheet}_模板.xlsx"
    return StreamingResponse(
        buffer,
        media_type="application/vnd.openxmlformats-officedocument.spreadsheetml.sheet",
        headers={"Content-Disposition": f"attachment; filename*=UTF-8''{quote(filename)}"},
    )


# ═══════════════════════════════════════════════════════════════════════════════
# 导出数据
# ═══════════════════════════════════════════════════════════════════════════════


@router.get("/{wp_id}/export-data")
async def n5_export_data(
    wp_id: str,
    sheet: str = Query("N5-5", description="Sheet编码: N5-2 / N5-5"),
    db: AsyncSession = Depends(get_db),
    current_user: User = Depends(get_current_user),
) -> StreamingResponse:
    """导出含当前数据的xlsx（纳税调整分sheet导出）

    Query params:
        sheet: 指定导出sheet（默认N5-5）
            - N5-2: 所得税费用明细表数据
            - N5-5: 纳税调整明细表（按分类分sheet导出，含数据）

    Requirements: 4.6
    """
    _validate_sheet(sheet)

    rows_data = await _load_rows_data(wp_id, sheet, db)

    if sheet == "N5-5":
        # N5-5 纳税调整：分类拆分为多sheet导出
        wb = _create_n5_5_split_template()
        # 按 category 字段分组写入各sheet
        cat_rows: dict[str, list[dict]] = {cat: [] for cat in _TAX_ADJUSTMENT_CATEGORIES}
        for row_data in rows_data:
            cat = row_data.get("category", "其他")
            if cat not in cat_rows:
                cat = "其他"
            cat_rows[cat].append(row_data)

        for cat_name in _TAX_ADJUSTMENT_CATEGORIES:
            ws = wb[cat_name[:31]]
            for row_idx, row_data in enumerate(cat_rows[cat_name], start=2):
                values = _export_row("N5-5", row_data)
                for col_idx, val in enumerate(values, 1):
                    ws.cell(row=row_idx, column=col_idx, value=val)
    else:
        # N5-2: 单sheet导出
        wb = _create_template_wb(sheet)
        ws = wb.active
        for row_idx, row_data in enumerate(rows_data, start=2):
            values = _export_row(sheet, row_data)
            for col_idx, val in enumerate(values, 1):
                ws.cell(row=row_idx, column=col_idx, value=val)

    buffer = io.BytesIO()
    wb.save(buffer)
    buffer.seek(0)

    filename = f"N5所得税费用_{sheet}_数据.xlsx"
    return StreamingResponse(
        buffer,
        media_type="application/vnd.openxmlformats-officedocument.spreadsheetml.sheet",
        headers={"Content-Disposition": f"attachment; filename*=UTF-8''{quote(filename)}"},
    )


# ═══════════════════════════════════════════════════════════════════════════════
# 导入数据
# ═══════════════════════════════════════════════════════════════════════════════


@router.post("/{wp_id}/import-data")
async def n5_import_data(
    wp_id: str,
    sheet: str = Query("N5-5", description="Sheet编码: N5-2 / N5-5"),
    file: UploadFile = File(...),
    db: AsyncSession = Depends(get_db),
    current_user: User = Depends(get_current_user),
) -> dict[str, Any]:
    """导入xlsx解析写入checklist_responses

    Query params:
        sheet: 指定导入目标sheet（默认N5-5）
            - N5-2: 所得税费用明细表
            - N5-5: 纳税调整明细表（支持分sheet或单sheet导入）

    验证：文件扩展名(.xlsx/.xls)、列头匹配、行数限制

    Returns:
        { imported_count: int, message: str, warnings: list[str] }

    Requirements: 4.6
    """
    _validate_sheet(sheet)

    # 文件扩展名校验
    filename = file.filename or ""
    suffix = Path(filename).suffix.lower()
    if suffix not in (".xlsx", ".xls"):
        raise HTTPException(400, f"不支持的文件类型: {suffix}，仅支持 .xlsx/.xls")

    content = await file.read()
    warnings: list[str] = []
    all_rows: list[dict] = []

    try:
        wb = load_workbook(io.BytesIO(content), data_only=True, read_only=True)
    except Exception:
        raise HTTPException(400, "无法解析xlsx文件，请检查文件格式")

    if sheet == "N5-5":
        # N5-5 支持多sheet导入（分类sheet）或单sheet
        for ws_name in wb.sheetnames:
            ws = wb[ws_name]
            actual_headers = _get_actual_headers(ws)
            # 检查是否包含N5-5的必要列头
            expected_cols = {"纳税调整项目", "调增金额", "调减金额"}
            if not expected_cols.issubset(set(actual_headers)):
                # 跳过非数据sheet
                continue

            missing = _validate_columns(ws, "N5-5")
            if missing:
                warnings.append(f"Sheet '{ws_name}' 缺少列: {missing}")
                continue

            row_count = 0
            for row_data in ws.iter_rows(min_row=2, values_only=True):
                # 跳过全空行
                if not any(cell is not None and str(cell).strip() for cell in row_data):
                    continue
                parsed = _parse_row("N5-5", row_data, actual_headers)
                # 用sheet名作为分类（如果行数据没有分类则填补）
                if not parsed.get("category") and ws_name in _TAX_ADJUSTMENT_CATEGORIES:
                    parsed["category"] = ws_name
                all_rows.append(parsed)
                row_count += 1
                if row_count >= _ROW_LIMIT:
                    warnings.append(f"Sheet '{ws_name}' 超过{_ROW_LIMIT}行限制，已截断")
                    break
    else:
        # N5-2: 单sheet导入
        ws = wb.active
        actual_headers = _get_actual_headers(ws)
        missing = _validate_columns(ws, sheet)
        if missing:
            wb.close()
            raise HTTPException(400, f"列头不匹配，缺少: {missing}")

        for row_data in ws.iter_rows(min_row=2, values_only=True):
            if not any(cell is not None and str(cell).strip() for cell in row_data):
                continue
            parsed = _parse_row(sheet, row_data, actual_headers)
            all_rows.append(parsed)
            if len(all_rows) >= _ROW_LIMIT:
                warnings.append(f"超过{_ROW_LIMIT}行限制，已截断")
                break

    wb.close()

    if not all_rows:
        raise HTTPException(400, "未从文件中解析到有效数据行")

    # 写入 checklist_responses
    item_id = _SHEET_ITEM_ID[sheet]
    rows_json = json.dumps(all_rows, ensure_ascii=False)

    # Upsert: 先检查是否存在
    existing = await db.execute(
        sa.text(
            "SELECT id FROM checklist_responses "
            "WHERE wp_id = :wp_id AND item_id = :item_id LIMIT 1"
        ),
        {"wp_id": wp_id, "item_id": item_id},
    )
    row = existing.fetchone()

    if row:
        await db.execute(
            sa.text(
                "UPDATE checklist_responses SET remark = :remark "
                "WHERE wp_id = :wp_id AND item_id = :item_id"
            ),
            {"remark": rows_json, "wp_id": wp_id, "item_id": item_id},
        )
    else:
        await db.execute(
            sa.text(
                "INSERT INTO checklist_responses (id, wp_id, item_id, remark) "
                "VALUES (:id, :wp_id, :item_id, :remark)"
            ),
            {
                "id": str(uuid4()),
                "wp_id": wp_id,
                "item_id": item_id,
                "remark": rows_json,
            },
        )

    await db.commit()

    result: dict[str, Any] = {
        "imported_count": len(all_rows),
        "message": f"成功导入 {len(all_rows)} 行到 {sheet}",
    }
    if warnings:
        result["warnings"] = warnings

    return result

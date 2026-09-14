"""F0 采购循环函证 — 导入导出（3 端点）.

端点：
- POST /api/workpapers/{wp_id}/f0/export-template?sheet=F0-5|F0-6
- POST /api/workpapers/{wp_id}/f0/export-data?sheet=F0-5|F0-6
- POST /api/workpapers/{wp_id}/f0/import-data?sheet=F0-5|F0-6

F0-5 预付账款替代程序：4 区块分 4 个 sheet（①期后收货/②期末余额证据/③本期付款/④本期采购入库）
F0-6 应付账款替代程序：4 区块分 4 个 sheet（①余额支持证据/②期后付款/③本期采购入库/④本期付款）

列定义与前端 blockColumnConfigsF05.ts / blockColumnConfigsF06.ts 对齐。
首列「被函证单位」用于导入时按供应商归组。
"""
from __future__ import annotations

import json
import logging
from typing import Any
from uuid import uuid4

import sqlalchemy as sa
from fastapi import APIRouter, Depends, File, HTTPException, Query, UploadFile
from fastapi.responses import StreamingResponse
from openpyxl import Workbook, load_workbook
from openpyxl.styles import Font
from openpyxl.utils import get_column_letter
from sqlalchemy.ext.asyncio import AsyncSession

from app.core.database import get_db
from app.deps import get_current_user
from app.models.core import User

from ._cycle_import_export_common import safe_float, safe_str, workbook_to_response

logger = logging.getLogger(__name__)

router = APIRouter(tags=["f0-import-export"])

ROW_LIMIT = 500

# ── sheet 编码 → parsed_data.html_data 中的 sheet 名 ──────────────────────────
_SHEET_NAME_MAP: dict[str, str] = {
    "F0-5": "替代程序F0-5",
    "F0-6": "替代程序F0-6",
}

# ── 凭证公共列 ────────────────────────────────────────────────────────────────
_VOUCHER_COLS: list[tuple[str, str, bool]] = [
    ("日期", "voucher_date", False),
    ("凭证编号", "voucher_no", False),
    ("业务内容", "business_desc", False),
    ("对方科目", "counter_account", False),
    ("金额", "voucher_amount", True),
]

# ── F0-5 四区块列定义（对齐 blockColumnConfigsF05.ts）─────────────────────────
_F05_BLOCKS: dict[str, dict[str, Any]] = {
    "block1": {
        "sheet_title": "①期后收货检查",
        "columns": [
            ("被函证单位", "_entity", False),
            ("序号", "seq", True),
            *_VOUCHER_COLS,
            ("入库单日期/编号", "inbound_date_no", False),
            ("品名", "inbound_product", False),
            ("单位", "inbound_unit", False),
            ("数量", "inbound_qty", True),
            ("发票日期/编号", "invoice_date_no", False),
            ("对手方", "invoice_counterparty", False),
            ("发票金额", "invoice_amount", True),
            ("索引号", "ref_index", False),
            ("是否异常", "is_abnormal", False),
        ],
    },
    "block2": {
        "sheet_title": "②期末余额支持性证据",
        "columns": [
            ("被函证单位", "_entity", False),
            ("序号", "seq", True),
            *_VOUCHER_COLS,
            ("付款审批日期/编号", "approval_date_no", False),
            ("是否恰当审批", "approval_ok", False),
            ("银行回单日期", "bank_date", False),
            ("收款方", "bank_payee", False),
            ("银行回单金额", "bank_amount", True),
            ("合同供应商", "contract_vendor", False),
            ("合同金额", "contract_amount", True),
            ("预付比例", "prepay_ratio", False),
            ("索引号", "ref_index", False),
            ("是否异常", "is_abnormal", False),
        ],
    },
    "block3": {
        "sheet_title": "③本期付款检查",
        "columns": [
            ("被函证单位", "_entity", False),
            ("序号", "seq", True),
            ("日期", "voucher_date", False),
            ("凭证编号", "voucher_no", False),
            ("业务内容", "business_desc", False),
            ("对方科目", "counter_account", False),
            ("金额", "payment_amount", True),
            ("付款审批日期/编号", "approval_date_no", False),
            ("是否恰当审批", "approval_ok", False),
            ("银行回单日期", "bank_date", False),
            ("收款方", "bank_payee", False),
            ("银行回单金额", "bank_amount", True),
            ("发票日期/编号", "invoice_date_no", False),
            ("对手方", "invoice_counterparty", False),
            ("发票金额", "invoice_amount", True),
            ("索引号", "ref_index", False),
            ("是否异常", "is_abnormal", False),
        ],
    },
    "block4": {
        "sheet_title": "④本期采购入库证据",
        "columns": [
            ("被函证单位", "_entity", False),
            ("序号", "seq", True),
            *_VOUCHER_COLS,
            ("入库单日期/编号", "inbound_date_no", False),
            ("品名", "inbound_product", False),
            ("单位", "inbound_unit", False),
            ("数量", "inbound_qty", True),
            ("合同日期/编号", "contract_date_no", False),
            ("供应商", "contract_vendor", False),
            ("合同金额", "contract_amount", True),
            ("发票日期/编号", "invoice_date_no", False),
            ("对手方", "invoice_counterparty", False),
            ("发票金额", "invoice_amount", True),
            ("索引号", "ref_index", False),
            ("是否异常", "is_abnormal", False),
        ],
    },
}

# ── F0-6 四区块列定义（对齐 blockColumnConfigsF06.ts）─────────────────────────
_F06_BLOCKS: dict[str, dict[str, Any]] = {
    "block1": {
        "sheet_title": "①应付余额支持性证据",
        "columns": [
            ("被函证单位", "_entity", False),
            ("序号", "seq", True),
            *_VOUCHER_COLS,
            ("入库单日期/编号", "inbound_date_no", False),
            ("品名", "inbound_product", False),
            ("单位", "inbound_unit", False),
            ("数量", "inbound_qty", True),
            ("合同日期", "contract_date", False),
            ("供应商", "contract_vendor", False),
            ("合同金额", "contract_amount", True),
            ("发票日期/编号", "invoice_date_no", False),
            ("对手方", "invoice_counterparty", False),
            ("发票金额", "invoice_amount", True),
            ("索引号", "ref_index", False),
            ("是否异常", "is_abnormal", False),
        ],
    },
    "block2": {
        "sheet_title": "②期后付款检查",
        "columns": [
            ("被函证单位", "_entity", False),
            ("序号", "seq", True),
            ("日期", "voucher_date", False),
            ("凭证编号", "voucher_no", False),
            ("业务内容", "business_desc", False),
            ("对方科目", "counter_account", False),
            ("金额", "payment_amount", True),
            ("付款审批日期/编号", "approval_date_no", False),
            ("是否恰当审批", "approval_ok", False),
            ("银行回单日期", "bank_date", False),
            ("收款方", "bank_payee", False),
            ("银行回单金额", "bank_amount", True),
            ("发票日期/编号", "invoice_date_no", False),
            ("对手方", "invoice_counterparty", False),
            ("发票金额", "invoice_amount", True),
            ("索引号", "ref_index", False),
            ("是否异常", "is_abnormal", False),
        ],
    },
    "block3": {
        "sheet_title": "③本期采购入库证据",
        "columns": [
            ("被函证单位", "_entity", False),
            ("序号", "seq", True),
            *_VOUCHER_COLS,
            ("入库单日期/编号", "inbound_date_no", False),
            ("品名", "inbound_product", False),
            ("单位", "inbound_unit", False),
            ("数量", "inbound_qty", True),
            ("合同日期/编号", "contract_date_no", False),
            ("供应商", "contract_vendor", False),
            ("合同金额", "contract_amount", True),
            ("发票日期/编号", "invoice_date_no", False),
            ("对手方", "invoice_counterparty", False),
            ("发票金额", "invoice_amount", True),
            ("索引号", "ref_index", False),
            ("是否异常", "is_abnormal", False),
        ],
    },
    "block4": {
        "sheet_title": "④本期付款检查",
        "columns": [
            ("被函证单位", "_entity", False),
            ("序号", "seq", True),
            ("日期", "voucher_date", False),
            ("凭证编号", "voucher_no", False),
            ("业务内容", "business_desc", False),
            ("对方科目", "counter_account", False),
            ("金额", "payment_amount", True),
            ("付款审批日期/编号", "approval_date_no", False),
            ("是否恰当审批", "approval_ok", False),
            ("银行回单日期", "bank_date", False),
            ("收款方", "bank_payee", False),
            ("银行回单金额", "bank_amount", True),
            ("发票日期/编号", "invoice_date_no", False),
            ("对手方", "invoice_counterparty", False),
            ("发票金额", "invoice_amount", True),
            ("索引号", "ref_index", False),
            ("是否异常", "is_abnormal", False),
        ],
    },
}

_F05_GUIDANCE = [
    "F0-5 预付账款及采购替代程序 导入说明",
    "",
    "本工作簿含 4 个区块 sheet（期后收货/期末余额证据/本期付款/本期采购入库）。",
    "「被函证单位」列用于将行归入对应供应商；同名单位的行会合并到同一项目下。",
    "付款检查比例、入库检查比例等由前端根据余额数据自动重算。",
]

_F06_GUIDANCE = [
    "F0-6 应付账款及采购替代程序 导入说明",
    "",
    "本工作簿含 4 个区块 sheet（余额支持证据/期后付款/本期采购入库/本期付款）。",
    "「被函证单位」列用于将行归入对应供应商；同名单位的行会合并到同一项目下。",
    "付款检查比例、入库检查比例等由前端根据余额数据自动重算。",
]


def _validate_sheet(sheet: str) -> None:
    if sheet not in _SHEET_NAME_MAP:
        raise HTTPException(400, f"不支持的sheet: {sheet}。支持: {sorted(_SHEET_NAME_MAP)}")


def _get_blocks(sheet: str) -> dict[str, dict[str, Any]]:
    return _F05_BLOCKS if sheet == "F0-5" else _F06_BLOCKS


def _get_guidance(sheet: str) -> list[str]:
    return _F05_GUIDANCE if sheet == "F0-5" else _F06_GUIDANCE


def _get_format_key(sheet: str) -> str:
    return "alternative-f05-v1" if sheet == "F0-5" else "alternative-f06-v1"


async def _load_html_data(db: AsyncSession, wp_id: str, sheet_name: str) -> dict:
    """读取 parsed_data.html_data[sheet_name]。"""
    result = await db.execute(
        sa.text(
            "SELECT parsed_data FROM working_paper WHERE id = :wp_id AND is_deleted = false"
        ),
        {"wp_id": wp_id},
    )
    row = result.fetchone()
    if not row or not row.parsed_data:
        return {}
    parsed = row.parsed_data
    if isinstance(parsed, str):
        try:
            parsed = json.loads(parsed)
        except (json.JSONDecodeError, TypeError):
            return {}
    html_data = (parsed or {}).get("html_data", {}) if isinstance(parsed, dict) else {}
    sheet_data = html_data.get(sheet_name)
    return sheet_data if isinstance(sheet_data, dict) else {}


async def _save_html_data(db: AsyncSession, wp_id: str, sheet_name: str, payload: dict) -> None:
    """将 payload 写回 parsed_data.html_data[sheet_name]。"""
    result = await db.execute(
        sa.text(
            "SELECT parsed_data FROM working_paper WHERE id = :wp_id AND is_deleted = false"
        ),
        {"wp_id": wp_id},
    )
    row = result.fetchone()
    if not row:
        raise HTTPException(404, f"底稿不存在: {wp_id}")
    parsed = row.parsed_data or {}
    if isinstance(parsed, str):
        try:
            parsed = json.loads(parsed)
        except (json.JSONDecodeError, TypeError):
            parsed = {}
    if not isinstance(parsed, dict):
        parsed = {}
    parsed.setdefault("html_data", {})
    parsed["html_data"][sheet_name] = payload
    await db.execute(
        sa.text(
            "UPDATE working_paper SET parsed_data = CAST(:pd AS jsonb), updated_at = NOW() "
            "WHERE id = :wp_id"
        ),
        {"pd": json.dumps(parsed, ensure_ascii=False), "wp_id": wp_id},
    )
    await db.commit()


# ── workbook 构建工具 ─────────────────────────────────────────────────────────

def _append_header_sheet(
    ws, columns: list[tuple[str, str, bool]], *, title: str | None = None
) -> int:
    """写标题 + 表头，返回表头所在行号。"""
    header_row = 1
    headers = [c[0] for c in columns]
    if title:
        ws.append([title])
        ws.merge_cells(start_row=1, start_column=1, end_row=1, end_column=max(len(headers), 1))
        ws["A1"].font = Font(bold=True, size=12)
        header_row = 2
    ws.append(headers)
    for h in ws[header_row]:
        h.font = Font(bold=True)
    ws.freeze_panes = f"A{header_row + 1}"
    for idx in range(1, len(headers) + 1):
        ws.column_dimensions[get_column_letter(idx)].width = 14
    return header_row


def _add_guidance(wb: Workbook, lines: list[str]) -> None:
    ws = wb.create_sheet("编制说明")
    for line in lines:
        ws.append([line])
    ws.column_dimensions["A"].width = 80


def _build_workbook(sheet: str, companies: list[dict] | None) -> Workbook:
    """构建 F0-5/F0-6 的多 sheet 工作簿。"""
    blocks = _get_blocks(sheet)
    guidance = _get_guidance(sheet)
    wb = Workbook()
    wb.remove(wb.active)
    for block_key, cfg in blocks.items():
        ws = wb.create_sheet(cfg["sheet_title"])
        columns = cfg["columns"]
        _append_header_sheet(ws, columns)
        for company in companies or []:
            entity = safe_str(company.get("entity_name"))
            for r in company.get(f"{block_key}_rows", []) or []:
                out_row = []
                for _header, key, is_num in columns:
                    if key == "_entity":
                        out_row.append(entity)
                    elif is_num:
                        out_row.append(safe_float(r.get(key)))
                    else:
                        out_row.append(safe_str(r.get(key)))
                ws.append(out_row)
    _add_guidance(wb, guidance)
    return wb


# ── xlsx 解析工具 ─────────────────────────────────────────────────────────────

def _content_io(content: bytes):
    import io
    return io.BytesIO(content)


def _parse_sheet_rows(
    content: bytes, ws_title: str | None, columns: list[tuple[str, str, bool]], header_row: int
) -> list[dict]:
    """按列头解析指定 sheet 的数据行。"""
    wb = load_workbook(_content_io(content), read_only=True, data_only=True)
    try:
        ws = wb[ws_title] if ws_title and ws_title in wb.sheetnames else wb.active
        if ws is None:
            return []
        header_cells = next(ws.iter_rows(min_row=header_row, max_row=header_row))
        actual = [str(c.value).strip() if c.value is not None else "" for c in header_cells]
        expected = [c[0] for c in columns]
        missing = [h for h in expected if h not in actual]
        if missing:
            raise ValueError(f"sheet「{ws.title}」缺少列: {', '.join(missing)}")
        rows: list[dict] = []
        for values in ws.iter_rows(min_row=header_row + 1, values_only=True):
            if values is None or all(v is None for v in values):
                continue
            if len(rows) >= ROW_LIMIT:
                break
            vals = list(values) + [None] * max(0, len(actual) - len(values))
            out: dict[str, Any] = {"_row_id": uuid4().hex[:12]}
            for header, key, is_num in columns:
                idx = actual.index(header)
                raw = vals[idx] if idx < len(vals) else None
                out[key] = safe_float(raw) if is_num else safe_str(raw)
            rows.append(out)
        return rows
    finally:
        wb.close()


# ── 端点 ──────────────────────────────────────────────────────────────────────

@router.post("/api/workpapers/{wp_id}/f0/export-template")
async def export_template(
    wp_id: str,
    sheet: str = Query(...),
    current_user: User = Depends(get_current_user),
) -> StreamingResponse:
    _validate_sheet(sheet)
    wb = _build_workbook(sheet, None)
    label = "F0-5预付账款替代程序" if sheet == "F0-5" else "F0-6应付账款替代程序"
    return workbook_to_response(wb, f"{label}_模板.xlsx")


@router.post("/api/workpapers/{wp_id}/f0/export-data")
async def export_data(
    wp_id: str,
    sheet: str = Query(...),
    db: AsyncSession = Depends(get_db),
    current_user: User = Depends(get_current_user),
) -> StreamingResponse:
    _validate_sheet(sheet)
    sheet_name = _SHEET_NAME_MAP[sheet]
    data = await _load_html_data(db, wp_id, sheet_name)
    companies = data.get("companies") if isinstance(data.get("companies"), list) else []
    wb = _build_workbook(sheet, companies)
    label = "F0-5预付账款替代程序" if sheet == "F0-5" else "F0-6应付账款替代程序"
    return workbook_to_response(wb, f"{label}_数据.xlsx")


@router.post("/api/workpapers/{wp_id}/f0/import-data")
async def import_data(
    wp_id: str,
    sheet: str = Query(...),
    file: UploadFile = File(...),
    db: AsyncSession = Depends(get_db),
    current_user: User = Depends(get_current_user),
) -> dict[str, Any]:
    _validate_sheet(sheet)
    if not file.filename or not file.filename.endswith(".xlsx"):
        raise HTTPException(400, "请上传 .xlsx 格式文件")
    content = await file.read()
    if len(content) > 10 * 1024 * 1024:
        raise HTTPException(400, "文件大小不能超过10MB")

    sheet_name = _SHEET_NAME_MAP[sheet]
    blocks = _get_blocks(sheet)
    format_key = _get_format_key(sheet)
    default_item = "预付账款" if sheet == "F0-5" else "应付账款"

    errors: list[str] = []
    entity_map: dict[str, dict[str, Any]] = {}
    total = 0
    for block_key, cfg in blocks.items():
        try:
            rows = _parse_sheet_rows(content, cfg["sheet_title"], cfg["columns"], header_row=1)
        except ValueError as e:
            errors.append(str(e))
            continue
        except Exception:
            raise HTTPException(400, "无法解析xlsx文件")
        for r in rows:
            entity = safe_str(r.pop("_entity", "")) or "未命名供应商"
            company = entity_map.setdefault(
                entity,
                {
                    "_company_id": uuid4().hex[:12],
                    "entity_name": entity,
                    "_source": "import",
                    "sampling": {},
                    "balance": {"item_name": default_item},
                    "block1_rows": [],
                    "block2_rows": [],
                    "block3_rows": [],
                    "block4_rows": [],
                    "conclusion": {},
                },
            )
            company[f"{block_key}_rows"].append(r)
            total += 1

    if errors and not entity_map:
        return {"ok": False, "errors": errors, "imported_count": 0}

    companies = list(entity_map.values())
    for i, c in enumerate(companies, start=1):
        c["seq"] = i
    payload = {"_format": format_key, "companies": companies}
    await _save_html_data(db, wp_id, sheet_name, payload)
    out: dict[str, Any] = {"ok": True, "imported_count": total, "errors": errors}
    return out

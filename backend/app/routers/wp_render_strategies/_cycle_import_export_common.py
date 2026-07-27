"""循环底稿导入导出 — 通用工具（D4 模式）."""

from __future__ import annotations

import io
import json
import re
from typing import Any, Callable
from uuid import uuid4

import sqlalchemy as sa
from fastapi import APIRouter, Depends, File, HTTPException, Query, UploadFile
from fastapi.responses import StreamingResponse
from openpyxl import Workbook, load_workbook
from openpyxl.styles import Font
from openpyxl.utils import get_column_letter
from sqlalchemy.ext.asyncio import AsyncSession

from app.core.database import get_db
from app.deps import get_current_user
from app.models.core import User

ROW_LIMIT = 500

_NUMERIC_KEY_RE = re.compile(r"^m\d+$")


def is_numeric_field_key(key: str) -> bool:
    if key in {"seq", "term", "days", "rate"}:
        return True
    if _NUMERIC_KEY_RE.match(key):
        return True
    if key.endswith(("Amt", "Qty", "Balance", "Total", "Rate", "Days", "Coef", "Share", "Pct", "Decrease")):
        return True
    # G1 交易性金融资产等强数值语义后缀（避免 sppiResult 之类枚举被误判，故不含 Result）
    if key.endswith(("Quantity", "Cost", "Value", "Gain", "Income", "Price", "Diff", "Fee")):
        return True
    if "Amount" in key or "Balance" in key or key.endswith("Occurrence"):
        return True
    # G1 专属数值字段（后缀无法覆盖：公允变动/调整分录/审定等）
    if key in {
        "quantity", "initialCost", "unitFairValue", "fairValueChange", "cumulativeFVChange",
        "fvChangeInPL", "disposalProceeds", "unadjusted", "aje", "rje", "adjusted", "variance",
        "debit", "credit", "dividendPerShare", "dividendIncome", "margin", "notionalAmount",
        "quoteValue", "bookValue", "marketValue", "countDayQuantity", "countDayAmount",
        "bookUnitFv", "level1Diff", "level2Result", "level2Diff", "level3Result", "level3Diff",
        "testedValue", "activeDiff",
    }:
        return True
    return key in {
        "faceValue", "termDays", "overdueDays", "concentration", "debitAmount", "creditAmount",
        "debitMovement", "creditMovement",
        "debit", "credit", "increase", "decrease", "interest", "accruedInterest", "bookInterest",
        "variance", "openingQty", "openingAmt", "increaseQty", "increaseAmt", "decreaseQty",
        "decreaseAmt", "agingLt1", "aging1to2", "aging2to3", "agingGt3", "expectedDays",
        "payableInterest", "openingAdjusted", "closingBalance", "agingTotal", "postPaymentAmt",
        "postPaymentAmount", "averagePaymentDays", "contractUnitPrice",
        "estimatedAmount", "voucherAmount", "reportPeriodAmount",
        "auditedBalance", "auditedAgingLt1", "auditedAging1to2", "auditedAging2to3", "auditedAgingGt3",
        "h1Total", "h1Share", "h1Avg", "h2Total", "yearTotal", "priorTotal", "changeAmt", "changeRate",
        # H1-12 折旧测算导入
        "originalCost", "accDepBegin", "bookAccDepEnd", "bookDepreciation", "bookMonthly",
        "usefulLife", "salvageRate", "impairmentBegin", "impairmentEnd", "impairmentProvision",
        "volCoef", "currentAmt", "priorAmt", "changeAmount", "changeRatePct", "expectedDiff",
        "hangDays", "hangAmount", "qtySold", "qtyCost", "qtyDiff", "qtyDiffRate", "openingStock",
        "production", "purchase", "availableQty", "closingStock", "theoreticalQty", "theoreticalDiff",
        "amount", "adjustAmount", "noteAmount", "discountAmount", "openingBalance", "sharePct",
        "currentIssued", "currentAccepted", "closingUnadjusted", "closingAdjusted",
        "depositRate", "depositAmount",
        # I3 商誉明细/减值滚动数值字段
        "costOpening", "costIncrease", "costEnding", "costUnadj", "costAje", "costAudited",
        "impOpening", "impIncrease", "impEnding", "impUnadj", "impAje", "impAudited",
        "goodwillNetValue", "goodwillOriginal", "goodwillB1", "minorityB2",
        "assetGroupCarrying", "cguBookValue", "fairValueLessCost", "valueInUse",
        "impairmentAmount", "goodwillImpairment", "consolidatedGwImpairment",
        "otherAssetImpairment", "entryGoodwillCalc", "mergerCost", "netAssetFairValue",
        "shareholding", "accImpairmentBegin", "accImpairmentEnd", "currentImpairment",
        "consideration", "counterpartyNetAsset", "minorityInterest",
        "goodwillAmount", "identifiableNetAssetFV",
        # F4-2 应付账款明细（源表A:AA）
        "openingUnadjusted", "openingAje", "openingRje", "currentDebit", "currentCredit",
        "entityReclassification", "unadjustedAgingLt1", "unadjustedAging1to2",
        "unadjustedAging2to3", "unadjustedAgingGt3", "closingAje", "closingRje",
        "auditedAgingLt1", "auditedAging1to2", "auditedAging2to3", "auditedAgingGt3",
        "subsequentPayment",
        "openingAdjustment", "closingAdjustment",
        "currentRevenue", "currentCost", "currentGross", "currentMargin", "priorRevenue", "priorCost",
        "priorGross", "priorMargin", "revenueChange", "revenueChangeRate", "costChange", "costChangeRate",
        "marginChange", "relatedRevenue", "costRate",
        # G14 信用减值损失明细
        "currentUnadjusted", "currentAdjustment", "openingProvision", "currentProvision",
        "currentReversal", "currentWriteoff", "otherMovement", "closingProvision",
        # G12 净敞口套期收益
        "instrumentOpeningFV", "instrumentClosingFV", "instrumentFVChange",
        "itemOpeningFV", "itemClosingFV", "itemFVChange",
        "hedgeRatio", "profitLossAmount", "ineffectiveness",
        "priorUnadjusted", "priorAdjustment",
        # G11 投资收益
        "currentIncome", "currentOpening", "currentClosing", "priorOpening", "priorClosing",
        "creditAmount",
        # G7 长期股权投资明细表 / 同控初始计量
        "holdingRatio", "votingRatio", "ownershipRatio", "shareholdingRatio",
        "purchaseRatio", "transactionNo",
        "openingInvestCost", "openingEquityAdj", "openingImpairment", "openingBookValue",
        "openingAuditedCost", "openingAuditedEquity", "openingAuditedImpairment", "openingAuditedNetValue",
        "increaseNewInvest", "increaseEquityMethod", "decreaseDisposal", "decreaseEquityAdj",
        "impairmentProvision", "impairmentReversal", "investeeNetProfit",
        "holdingRatioChange", "otherComprehensiveIncome", "otherEquityChange", "profitDistribution",
        "closingInvestCost", "closingEquityAdj", "closingSubtotal",
        "closingImpairment", "closingBookValue", "auditAdjustment", "auditedAmount",
        "recoverableAmount", "disposalGainLoss",
        "investeeNetAssets", "shareOfNetAssets", "goodwill",
        "internalTransElim", "unrecognizedLoss", "equityMethodIncome",
        "currentOCI", "dividendIncome",
        "ownerEquityBookValue", "initialInvestmentCost", "cashConsideration",
        "nonCashAssetBookValue", "debtBookValue", "equitySecuritiesFaceValue",
        "contingentConsideration", "totalConsideration",
        "capitalReserveRetainedEarningsAdjustment", "consideration",
        "netAssetsBookValue", "priorInvestmentAdjustments",
        # I1 无形资产 / H1 固定资产共通数值字段
        "costOpening", "costIncrease", "costDecrease", "costClosing",
        "amortOpening", "amortProvision", "amortTransferOut", "amortClosing",
        "depOpening", "depProvision", "depReversal", "depClosing",
        "impairOpening", "impairProvision", "impairReversal", "impairClosing",
        "accAmortization", "impairment", "difference", "audited", "unadjusted", "aje", "rje",
        "usefulLife", "salvageRate", "entryAmount", "originalCost",
        "amortTotal", "adminExpense", "salesExpense", "mfgExpense", "rdExpense", "otherExpense",
        "allocTotal", "allocRatio",
        # F2-43/F2-44 分配基数
        "allocationBase",
        # F2-47/F2-48 数量与库龄四档
        "qty", "within1y", "y1to2", "y2to3", "over3y",
        # K5 预计负债（K5-4 质保测算 / K5-5 弃置现值 / K5-6 诉讼）
        "revenue", "estimatedExpense", "bookProvision", "periodProvision", "periodUsed",
        "futureExpense", "expectedYears", "periodIncrease", "interestAdjustment", "estimatedLoss",
        # K11-2 减值准备变动
        "allowanceOpening", "allowanceWriteoff", "allowanceEnding",
    }


def safe_float(val: Any) -> float:
    if val is None:
        return 0.0
    try:
        return float(val)
    except (ValueError, TypeError):
        return 0.0


def safe_str(val: Any) -> str:
    return "" if val is None else str(val).strip()


def col_val(row: tuple, headers: list[str], name: str) -> Any:
    try:
        idx = headers.index(name)
        return row[idx] if idx < len(row) else None
    except ValueError:
        return None


def apply_header_aliases(
    actual_headers: list[str],
    header_aliases: dict[str, list[str]] | None,
) -> list[str]:
    """将实际表头中的旧名归一为规范名。header_aliases: {规范名: [别名,…]}"""
    if not header_aliases:
        return list(actual_headers)
    alias_to_canonical: dict[str, str] = {}
    for canonical, aliases in header_aliases.items():
        for a in aliases or []:
            if a:
                alias_to_canonical[str(a).strip()] = canonical
    return [alias_to_canonical.get(h, h) for h in actual_headers]


def parse_row_by_headers(
    row: tuple,
    headers: list[str],
    field_keys: list[str],
    *,
    expected_headers: list[str] | None = None,
) -> dict[str, Any]:
    """按表头解析一行。

    - 若提供 expected_headers：按名称匹配（与列顺序无关，缺列填空）
    - 否则：按位置 zip（兼容旧调用）
    """
    values = list(row) + [None] * max(0, len(headers) - len(row))
    out: dict[str, Any] = {"id": str(uuid4())}
    if expected_headers is not None:
        for h, key in zip(expected_headers, field_keys):
            raw = col_val(tuple(values), headers, h)
            if is_numeric_field_key(key):
                out[key] = safe_float(raw)
            else:
                out[key] = safe_str(raw)
        return out
    for h, key in zip(headers, field_keys):
        idx = headers.index(h)
        raw = values[idx] if idx < len(values) else None
        if is_numeric_field_key(key):
            out[key] = safe_float(raw)
        else:
            out[key] = safe_str(raw)
    return out


def recalculate_computed_fields(rows: list[dict], computed_fields: dict[str, str] | None) -> list[dict]:
    """Recalculate computed columns before export.

    computed_fields format: {"target_field": "formula"}
    Supported formulas: "a + b + ...", "a - b", "a * b", "a / b" (simple two-operand arithmetic)
    """
    if not computed_fields:
        return rows
    for row in rows:
        for target, formula in computed_fields.items():
            try:
                formula = formula.strip()
                if ' + ' in formula:
                    parts = formula.split(' + ')
                    val = sum(float(row.get(p.strip(), 0) or 0) for p in parts)
                elif ' - ' in formula:
                    parts = formula.split(' - ', 1)
                    val = float(row.get(parts[0].strip(), 0) or 0) - float(row.get(parts[1].strip(), 0) or 0)
                elif ' * ' in formula:
                    parts = formula.split(' * ', 1)
                    val = float(row.get(parts[0].strip(), 0) or 0) * float(row.get(parts[1].strip(), 0) or 0)
                elif ' / ' in formula:
                    parts = formula.split(' / ', 1)
                    divisor = float(row.get(parts[1].strip(), 0) or 0)
                    if divisor == 0:
                        continue  # skip division by zero
                    val = float(row.get(parts[0].strip(), 0) or 0) / divisor
                else:
                    continue
                row[target] = round(val, 2)
            except (ValueError, TypeError, ZeroDivisionError):
                pass
    return rows


def export_row_by_keys(data: dict, field_keys: list[str]) -> list:
    row: list[Any] = []
    for key in field_keys:
        val = data.get(key)
        if val is None:
            row.append("")
        elif isinstance(val, float) or isinstance(val, int):
            row.append(round(val, 2))
        else:
            row.append(val)
    return row


async def load_json_payload(
    db: AsyncSession,
    wp_id: str,
    item_id: str,
    field: str = "conclusion",
) -> Any:
    """读取 checklist_responses 字段原始 JSON（list / dict / 其它）。"""
    result = await db.execute(
        sa.text(f"SELECT {field} FROM checklist_responses WHERE wp_id = :wp_id AND item_id = :iid LIMIT 1"),
        {"wp_id": wp_id, "iid": item_id},
    )
    row = result.fetchone()
    if not row or not getattr(row, field, None):
        return None
    try:
        return json.loads(getattr(row, field))
    except (json.JSONDecodeError, TypeError):
        return None


def keyed_store_to_rows(store: dict[str, Any], keyed_by: str = "rowKey") -> list[dict]:
    """将 {rowKey: {fields...}} 转为导出行列表。"""
    rows: list[dict] = []
    for key, fields in store.items():
        if not isinstance(fields, dict):
            continue
        row = {**fields, keyed_by: key}
        if keyed_by not in fields:
            row[keyed_by] = key
        rows.append(row)
    return rows


def rows_to_keyed_store(
    rows: list[dict],
    keyed_by: str = "rowKey",
    keep_keys: list[str] | None = None,
) -> dict[str, dict]:
    """将导出行列表写回 {rowKey: {fields...}}。"""
    store: dict[str, dict] = {}
    for row in rows:
        key = safe_str(row.get(keyed_by))
        if not key:
            continue
        if keep_keys is not None:
            cell = {k: row.get(k) for k in keep_keys if k in row and k != keyed_by}
        else:
            cell = {k: v for k, v in row.items() if k not in {keyed_by, "id"}}
        store[key] = cell
    return store


async def load_json_rows(
    db: AsyncSession,
    wp_id: str,
    item_id: str,
    field: str = "conclusion",
    *,
    keyed_by: str | None = None,
) -> list[dict]:
    parsed = await load_json_payload(db, wp_id, item_id, field=field)
    if parsed is None:
        return []
    if isinstance(parsed, list):
        return parsed
    if isinstance(parsed, dict):
        # 兼容页面存 {rows:[...], conclusion:...} / {data:[...]} / {items:[...]}
        for key in ("rows", "data", "items"):
            candidate = parsed.get(key)
            if isinstance(candidate, list):
                return [row for row in candidate if isinstance(row, dict)]
        if keyed_by:
            return keyed_store_to_rows(parsed, keyed_by=keyed_by)
    return []


async def upsert_json_payload(
    db: AsyncSession,
    wp_id: str,
    item_id: str,
    payload: Any,
    field: str = "conclusion",
) -> None:
    """写入 checklist_responses 字段任意 JSON（list / dict）。"""
    proj = await db.execute(
        sa.text(
            "SELECT project_id FROM working_paper WHERE id = :wp_id AND is_deleted = false"
        ),
        {"wp_id": wp_id},
    )
    project_id = proj.scalar_one_or_none()
    if not project_id:
        raise ValueError(f"working_paper not found: {wp_id}")

    await db.execute(
        sa.text(f"""
            INSERT INTO checklist_responses (id, project_id, wp_id, item_id, {field}, updated_at, created_at)
            VALUES (:id, :project_id, :wp_id, :item_id, :payload, NOW(), NOW())
            ON CONFLICT (wp_id, item_id)
            DO UPDATE SET {field} = :payload, updated_at = NOW()
        """),
        {
            "id": str(uuid4()),
            "project_id": str(project_id),
            "wp_id": wp_id,
            "item_id": item_id,
            "payload": json.dumps(payload, ensure_ascii=False),
        },
    )
    await db.commit()


async def upsert_json_rows(
    db: AsyncSession,
    wp_id: str,
    item_id: str,
    rows_data: list[dict],
    field: str = "conclusion",
    *,
    keyed_by: str | None = None,
    keep_keys: list[str] | None = None,
) -> None:
    payload: Any = rows_data
    if keyed_by:
        payload = rows_to_keyed_store(rows_data, keyed_by=keyed_by, keep_keys=keep_keys)
    await upsert_json_payload(db, wp_id, item_id, payload, field=field)


def build_workbook_template(
    sheet_code: str,
    headers: list[str],
    *,
    title: str | None = None,
    subtitle: str | None = None,
    guidance: list[str] | None = None,
    prefill_rows: list[list[Any]] | None = None,
) -> Workbook:
    wb = Workbook()
    ws = wb.active
    ws.title = sheet_code
    header_row = 1
    if title:
        ws.append([title])
        ws.merge_cells(start_row=1, start_column=1, end_row=1, end_column=max(len(headers), 1))
        ws["A1"].font = Font(bold=True, size=12)
        header_row += 1
    if subtitle:
        ws.append([subtitle])
        header_row += 1
    ws.append(headers)
    ws.freeze_panes = f"A{header_row + 1}"
    for col_idx in range(1, len(headers) + 1):
        ws.column_dimensions[get_column_letter(col_idx)].width = 14
    if prefill_rows:
        for r in prefill_rows:
            ws.append(r)
    if guidance:
        ws_guide = wb.create_sheet("编制说明")
        ws_guide.append(["编制说明"])
        ws_guide.append([])
        for line in guidance:
            ws_guide.append([line])
        ws_guide.column_dimensions["A"].width = 80
    return wb


# ─── 动态账龄列头支持（aging-config-enhancement Task 12.1） ───────────────────

# 账龄期间标签（与前端 AGING_EXPORT_PERIOD_LABELS 保持一致）
AGING_PERIOD_LABELS: dict[str, str] = {
    "prior": "期初",
    "current": "期末未审",
    "audited": "期末审定",
}

# 三期科目（含期末未审 current）：D2/K1/K3/G5/F1；两期科目（仅期初/期末审定）：D3
_THREE_PERIOD_SUBJECTS = frozenset({"D2", "K1", "K3", "G5", "F1"})

# 期末两期科目（期末未审 + 期末审定，**无期初账龄**）：F4 应付账款
# （F4 源表 N:Q 为期末未审账龄、U:X 为期末审定账龄，落默认 [prior, audited] 会产错列头）
_CLOSING_TWO_PERIOD_SUBJECTS = frozenset({"F4"})


# 期间 key → DetailRow 中嵌套 aging 数据字段名
_AGING_PERIOD_FIELD: dict[str, str] = {
    "prior": "agingPrior",
    "current": "agingCurrent",
    "audited": "agingAudited",
}


def subject_aging_periods(subject: str) -> list[str]:
    """返回科目对应的账龄期间列表（有序）。

    - 三期科目（D2/K1/K3/G5/F1）：[prior, current, audited]
    - 期末两期科目（F4，无期初账龄）：[current, audited]
    - 其余两期科目（D3）：[prior, audited]
    """
    if subject in _THREE_PERIOD_SUBJECTS:
        return ["prior", "current", "audited"]
    if subject in _CLOSING_TWO_PERIOD_SUBJECTS:
        return ["current", "audited"]
    return ["prior", "audited"]


def build_aging_headers(segments: list[Any], periods: list[str]) -> list[str]:
    """按 segments 顺序 + 期间生成动态账龄列头（period-major）。

    格式：`{label}({period_label})`（期初/期末未审/期末审定），委托 _aging_export_headers 的
    header_fn 工厂，保证与导入端（12.2）列头匹配逻辑使用同一套格式。
    """
    from ._aging_export_headers import suffix_header_fn

    headers: list[str] = []
    for period in periods:
        header_fn = suffix_header_fn(f"({AGING_PERIOD_LABELS.get(period, period)})")
        headers.extend(header_fn(seg) for seg in segments)
    return headers


def aging_export_values(data: dict, segments: list[Any], periods: list[str]) -> list[Any]:
    """按 build_aging_headers 相同顺序，从嵌套 aging 结构提取导出值。

    读取 data['agingPrior'/'agingCurrent'/'agingAudited'][seg.key]（缺失取 0）。
    """
    from ._aging_export_headers import get_segment_keys

    keys = get_segment_keys(segments)
    values: list[Any] = []
    for period in periods:
        bucket = data.get(_AGING_PERIOD_FIELD.get(period, "")) or {}
        for key in keys:
            values.append(safe_float(bucket.get(key)) if isinstance(bucket, dict) else 0.0)
    return values


def match_import_aging(
    get_value: Callable[[str], Any],
    actual_headers: list[str],
    segments: list[Any],
    periods: list[str],
) -> tuple[dict[str, dict[str, float]], list[str]]:
    """按 build_aging_headers 相同格式，从导入行匹配账龄列 → nested keyed 结构。

    - 列头格式 `{label}({period_label})`（期初/期末未审/期末审定），与导出端（build_aging_headers）
      使用同一套格式，保证 export→import 往返一致。
    - 每段 label 命中列头 → 值写入 row[period_field][seg.key]；缺列（配置变更后旧模板缺该段）→ 初始化 0。
    - unmatched：actual_headers 中"看起来像账龄"但不属于当前 segments×periods 的列 → 返回供上层报 warning + 跳过。

    返回 (aging_nested, unmatched_headers)：
      aging_nested = {"agingPrior": {seg_key: float}, "agingCurrent": {...}, "agingAudited": {...}}

    Requirements: 8.2 (label 匹配) / 8.3 (unmatched warning+skip) / 8.4 (旧模板按 label 映射)
    """
    from ._aging_export_headers import _looks_like_aging

    expected: set[str] = set()
    aging: dict[str, dict[str, float]] = {}
    for period in periods:
        field = _AGING_PERIOD_FIELD.get(period, period)
        plabel = AGING_PERIOD_LABELS.get(period, period)
        bucket: dict[str, float] = {}
        for seg in segments:
            header = f"{seg.label}({plabel})"
            expected.add(header)
            bucket[seg.key] = safe_float(get_value(header))
        aging[field] = bucket

    unmatched = [
        h for h in actual_headers
        if h and h not in expected and _looks_like_aging(h)
    ]
    return aging, unmatched


async def resolve_aging_segments(db: AsyncSession, wp_id: str, subject: str) -> list[Any]:
    """解析某底稿对应项目 + 科目的有效账龄段列表。

    委托 _aging_export_headers.get_project_segments_from_wp（单一 DB 解析入口）。
    异常时回退到该科目默认预设段，保证导出不崩。
    """
    from ._aging_export_headers import get_project_segments_from_wp

    try:
        return await get_project_segments_from_wp(db, wp_id, subject)
    except Exception:  # noqa: BLE001 — 兜底：任何异常回退默认预设
        from app.services import aging_config_service as _acs

        default_preset = _acs.DEFAULT_SUBJECT_PRESETS.get(subject, _acs.AgingPreset.FIVE_YEAR)
        return _acs.resolve_segments(default_preset, None)


def workbook_to_response(wb: Workbook, filename: str) -> StreamingResponse:
    buf = io.BytesIO()
    wb.save(buf)
    buf.seek(0)
    from urllib.parse import quote
    fn = quote(filename)
    return StreamingResponse(
        buf,
        media_type="application/vnd.openxmlformats-officedocument.spreadsheetml.sheet",
        headers={"Content-Disposition": f"attachment; filename*=UTF-8''{fn}"},
    )


def parse_upload_xlsx(
    content: bytes,
    expected_headers: list[str],
    header_row: int = 1,
    *,
    header_aliases: dict[str, list[str]] | None = None,
    require_all_headers: bool = True,
) -> tuple[list[str], list[tuple]]:
    wb = load_workbook(io.BytesIO(content), read_only=True, data_only=True)
    ws = wb.active
    if ws is None:
        wb.close()
        raise ValueError("xlsx文件中无活动工作表")
    actual = [
        str(c.value).strip() if c.value else ""
        for c in next(ws.iter_rows(min_row=header_row, max_row=header_row))
    ]
    actual = apply_header_aliases(actual, header_aliases)
    if require_all_headers:
        missing = [h for h in expected_headers if h not in actual]
        if missing:
            wb.close()
            raise ValueError(f"缺少列: {', '.join(missing)}")
    else:
        # 至少要命中一张主键列或期望列交集，避免空表/错表被静默导入
        if not any(h in actual for h in expected_headers):
            wb.close()
            raise ValueError("未匹配到任何预期列，请检查是否为正确模板")
    rows: list[tuple] = []
    for row in ws.iter_rows(min_row=header_row + 1, values_only=True):
        if all(v is None for v in row):
            continue
        rows.append(row)
    wb.close()
    return actual, rows


def import_rows_generic(
    rows: list[tuple],
    headers: list[str],
    field_keys: list[str],
    parse_fn: Callable[[tuple, list[str]], dict] | None = None,
) -> tuple[list[dict], bool]:
    parse = parse_fn or (lambda r, h: parse_row_by_headers(r, h, field_keys))
    out: list[dict] = []
    truncated = False
    for i, row in enumerate(rows, start=1):
        if i > ROW_LIMIT:
            truncated = True
            break
        out.append(parse(row, headers))
    return out, truncated


def create_cycle_import_export_router(
    *,
    tag: str,
    api_prefix: str,
    specs: dict[str, dict[str, Any]],
    storage_field: str = "conclusion",
    header_row: int = 2,
) -> APIRouter:
    """根据 sheet 规格表生成三端点导入导出 router（D4 模式）."""
    router = APIRouter(tags=[tag])
    supported = set(specs.keys())

    def _validate(sheet: str) -> None:
        if sheet not in supported:
            raise HTTPException(400, f"不支持的sheet: {sheet}。支持: {sorted(supported)}")

    def _spec(sheet: str) -> dict[str, Any]:
        return specs[sheet]

    async def _load_payload_fallback(
        db: AsyncSession,
        wp_id: str,
        item_id: str,
        field: str,
    ) -> Any:
        payload = await load_json_payload(db, wp_id, item_id, field=field)
        if payload is not None:
            return payload
        other = "remark" if field == "conclusion" else "conclusion"
        return await load_json_payload(db, wp_id, item_id, field=other)

    async def _load_rows_fallback(
        db: AsyncSession,
        wp_id: str,
        item_id: str,
        field: str,
        keyed_by: str | None,
    ) -> list[dict]:
        rows = await load_json_rows(db, wp_id, item_id, field=field, keyed_by=keyed_by)
        if rows:
            return rows
        other = "remark" if field == "conclusion" else "conclusion"
        return await load_json_rows(db, wp_id, item_id, field=other, keyed_by=keyed_by)

    async def _upsert_payload_maybe_dual(
        db: AsyncSession,
        wp_id: str,
        item_id: str,
        payload: Any,
        field: str,
        *,
        dual: bool,
    ) -> None:
        await upsert_json_payload(db, wp_id, item_id, payload, field=field)
        if dual:
            other = "remark" if field == "conclusion" else "conclusion"
            await upsert_json_payload(db, wp_id, item_id, payload, field=other)

    async def _upsert_rows_maybe_dual(
        db: AsyncSession,
        wp_id: str,
        item_id: str,
        rows: list[dict],
        field: str,
        *,
        dual: bool,
        keyed_by: str | None = None,
        keep_keys: list[str] | None = None,
    ) -> None:
        await upsert_json_rows(
            db, wp_id, item_id, rows, field=field, keyed_by=keyed_by, keep_keys=keep_keys,
        )
        if dual:
            other = "remark" if field == "conclusion" else "conclusion"
            await upsert_json_rows(
                db, wp_id, item_id, rows, field=other, keyed_by=keyed_by, keep_keys=keep_keys,
            )

    async def _export_headers(sp: dict[str, Any], wp_id: str, db: AsyncSession) -> list[str]:
        """构建导出列头。含 aging 描述符时拼接动态账龄列头（Task 12.1）。"""
        aging = sp.get("aging")
        if not aging:
            return sp["headers"]
        subject = aging["subject"]
        periods = aging.get("periods") or subject_aging_periods(subject)
        segments = await resolve_aging_segments(db, wp_id, subject)
        return list(aging["base_headers"]) + build_aging_headers(segments, periods)

    async def _export_row(sp: dict[str, Any], d: dict, wp_id: str, db: AsyncSession, _seg_cache: dict) -> list[Any]:
        """构建单行导出值。含 aging 描述符时拼接动态账龄值（Task 12.1）。"""
        aging = sp.get("aging")
        if not aging:
            return export_row_by_keys(d, sp["field_keys"])
        subject = aging["subject"]
        periods = aging.get("periods") or subject_aging_periods(subject)
        segments = _seg_cache.get(subject)
        if segments is None:
            segments = await resolve_aging_segments(db, wp_id, subject)
            _seg_cache[subject] = segments
        base = export_row_by_keys(d, aging["base_field_keys"])
        return base + aging_export_values(d, segments, periods)

    @router.post(f"/api/workpapers/{{wp_id}}/{api_prefix}/export-template")
    async def export_template(
        wp_id: str,
        sheet: str = Query(...),
        db: AsyncSession = Depends(get_db),
        current_user: User = Depends(get_current_user),
    ) -> StreamingResponse:
        _validate(sheet)
        sp = _spec(sheet)
        build_wb = sp.get("build_workbook")
        if callable(build_wb):
            wb = build_wb(None, template_only=True)
            return workbook_to_response(wb, f"{sheet}_模板.xlsx")
        headers = await _export_headers(sp, wp_id, db)
        wb = build_workbook_template(
            sheet,
            headers,
            title=sp.get("title"),
            subtitle=sp.get("subtitle"),
            guidance=sp.get("guidance"),
            prefill_rows=sp.get("template_prefill"),
        )
        return workbook_to_response(wb, f"{sheet}_模板.xlsx")

    @router.post(f"/api/workpapers/{{wp_id}}/{api_prefix}/export-data")
    async def export_data(
        wp_id: str,
        sheet: str = Query(...),
        db: AsyncSession = Depends(get_db),
        current_user: User = Depends(get_current_user),
    ) -> StreamingResponse:
        _validate(sheet)
        sp = _spec(sheet)
        item_id = sp.get("item_id", f"{sheet}-rows")
        field = sp.get("storage_field", storage_field)
        build_wb = sp.get("build_workbook")
        if callable(build_wb):
            export_loader = sp.get("export_loader")
            if callable(export_loader):
                payload = await export_loader(db, wp_id)
            else:
                payload = await _load_payload_fallback(db, wp_id, item_id, field)
            wb = build_wb(payload, template_only=False)
            return workbook_to_response(wb, f"{sheet}_数据.xlsx")
        keyed_by = sp.get("keyed_by")
        item_ids = list(sp.get("item_id_candidates") or [sp.get("item_id", f"{sheet}-rows")])
        rows: list = []
        for iid in item_ids:
            rows = await _load_rows_fallback(db, wp_id, iid, field, keyed_by)
            if rows:
                break
        headers = await _export_headers(sp, wp_id, db)
        prefill = sp.get("template_prefill")
        wb = build_workbook_template(
            sheet,
            headers,
            title=sp.get("title"),
            subtitle=sp.get("subtitle"),
            guidance=sp.get("guidance"),
            prefill_rows=None if rows else prefill,
        )
        ws = wb[sheet]
        _seg_cache: dict = {}
        export_rows = rows if rows else []
        # Recalculate computed fields before export (Improvement #4)
        computed = sp.get("computed_fields")
        if computed:
            export_rows = recalculate_computed_fields(export_rows, computed)
        for d in export_rows:
            ws.append(await _export_row(sp, d, wp_id, db, _seg_cache))
        return workbook_to_response(wb, f"{sheet}_数据.xlsx")

    @router.post(f"/api/workpapers/{{wp_id}}/{api_prefix}/import-data")
    async def import_data(
        wp_id: str,
        sheet: str = Query(...),
        file: UploadFile = File(...),
        db: AsyncSession = Depends(get_db),
        current_user: User = Depends(get_current_user),
    ) -> dict[str, Any]:
        _validate(sheet)
        if not file.filename or not file.filename.endswith(".xlsx"):
            raise HTTPException(400, "请上传 .xlsx 格式文件")
        content = await file.read()
        if len(content) > 10 * 1024 * 1024:
            raise HTTPException(400, "文件大小不能超过10MB")
        sp = _spec(sheet)
        parse_import = sp.get("parse_import")
        if callable(parse_import):
            try:
                payload, errors, imported_count = parse_import(content)
            except ValueError as e:
                return {"ok": False, "errors": [str(e)], "imported_count": 0}
            except Exception:
                raise HTTPException(400, "无法解析xlsx文件")
            if errors and not imported_count:
                return {"ok": False, "errors": errors, "imported_count": 0}
            import_handler = sp.get("import_handler")
            if callable(import_handler):
                await import_handler(db, wp_id, payload)
            else:
                item_id = sp.get("item_id", f"{sheet}-rows")
                field = sp.get("storage_field", storage_field)
                await _upsert_payload_maybe_dual(
                    db, wp_id, item_id, payload, field,                 dual=bool(sp.get("dual_write")),
                )
            out_custom: dict[str, Any] = {
                "ok": True,
                "imported_count": imported_count,
                "errors": errors,
            }
            if imported_count >= ROW_LIMIT:
                out_custom["warning"] = f"数据行数超过{ROW_LIMIT}行限制，已截断"
            return out_custom

        hr = sp.get("header_row", header_row)
        aging = sp.get("aging")

        # 含 aging 描述符（K1-2/K3-2/G5-2）：按项目账龄配置动态 label 匹配（Task 12.2）
        if aging:
            subject = aging["subject"]
            periods = aging.get("periods") or subject_aging_periods(subject)
            segments = await resolve_aging_segments(db, wp_id, subject)
            base_headers = list(aging["base_headers"])
            base_field_keys = list(aging["base_field_keys"])
            try:
                actual, raw = parse_upload_xlsx(content, base_headers, header_row=hr)
            except ValueError as e:
                return {"ok": False, "errors": [str(e)], "imported_count": 0}
            except Exception:
                raise HTTPException(400, "无法解析xlsx文件")

            # unmatched 账龄列检测（一次性，基于实际列头）
            _, skipped = match_import_aging(lambda _h: None, actual, segments, periods)

            def _parse_aging(r: tuple, h: list[str]) -> dict:
                out_row: dict[str, Any] = {"id": str(uuid4())}
                for header, key in zip(base_headers, base_field_keys):
                    raw_v = col_val(r, h, header)
                    out_row[key] = safe_float(raw_v) if is_numeric_field_key(key) else safe_str(raw_v)
                aging_nested, _u = match_import_aging(
                    lambda x: col_val(r, h, x), h, segments, periods,
                )
                out_row.update(aging_nested)
                return out_row

            rows, truncated = import_rows_generic(raw, actual, base_field_keys, parse_fn=_parse_aging)
            item_id = sp.get("item_id", f"{sheet}-rows")
            field = sp.get("storage_field", storage_field)
            await _upsert_rows_maybe_dual(
                db, wp_id, item_id, rows, field,
                dual=bool(sp.get("dual_write")),
                keyed_by=sp.get("keyed_by"),
                keep_keys=sp.get("keep_keys"),
            )
            out: dict[str, Any] = {"ok": True, "imported_count": len(rows), "errors": []}
            if truncated:
                out["warning"] = f"数据行数超过{ROW_LIMIT}行限制，已截断"
            if skipped:
                out["skipped_columns"] = skipped
                out["warnings"] = [f"以下账龄列未匹配当前账龄配置，已跳过: {', '.join(skipped)}"]
            return out

        try:
            actual, raw = parse_upload_xlsx(
                content,
                sp["headers"],
                header_row=hr,
                header_aliases=sp.get("header_aliases"),
                require_all_headers=not bool(sp.get("allow_missing_headers")),
            )
        except ValueError as e:
            return {"ok": False, "errors": [str(e)], "imported_count": 0}
        except Exception:
            raise HTTPException(400, "无法解析xlsx文件")
        keys = sp["field_keys"]
        expected = list(sp["headers"])
        rows, truncated = import_rows_generic(
            raw,
            actual,
            keys,
            parse_fn=lambda r, h: parse_row_by_headers(
                r, h, keys, expected_headers=expected,
            ),
        )
        item_id = sp.get("item_id", f"{sheet}-rows")
        field = sp.get("storage_field", storage_field)
        await _upsert_rows_maybe_dual(
            db, wp_id, item_id, rows, field,
            dual=bool(sp.get("dual_write")),
            keyed_by=sp.get("keyed_by"),
            keep_keys=sp.get("keep_keys"),
        )
        for mid in sp.get("mirror_item_ids") or []:
            if mid and mid != item_id:
                await _upsert_rows_maybe_dual(
                    db, wp_id, mid, rows, field,
                    dual=bool(sp.get("dual_write")),
                    keyed_by=sp.get("keyed_by"),
                    keep_keys=sp.get("keep_keys"),
                )
        out2: dict[str, Any] = {"ok": True, "imported_count": len(rows), "errors": []}
        if truncated:
            out2["warning"] = f"数据行数超过{ROW_LIMIT}行限制，已截断"
        return out2

    return router


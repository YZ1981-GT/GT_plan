"""E1 货币资金 — 银行对账单/流水 OCR 识别端点（E1-31）

POST /api/workpapers/{wp_id}/e1/statement-ocr
Content-Type: multipart/form-data
Body: file (PDF/image)

OCR + LLM 提取流水明细行，供前端确认后回填 E1-31 流水库并支持双向匹配。
"""
from __future__ import annotations

import json
import logging
from pathlib import Path
from typing import Any
from uuid import UUID, uuid4

import aiofiles
from fastapi import APIRouter, Depends, File, Form, HTTPException, UploadFile
from pydantic import BaseModel
from sqlalchemy.ext.asyncio import AsyncSession

from app.core.database import get_db
from app.deps import get_current_user
from app.models.core import User
from app.services.attachment_ocr_writeback import (
    finalize_linked_ocr_writeback,
    load_reusable_ocr,
    parse_optional_uuid,
    provision_wp_linked_attachment,
)
from app.services.attachment_service import AttachmentService
from app.services.llm_client import chat_completion
from app.services.unified_ocr_service import UnifiedOCRService

logger = logging.getLogger(__name__)

router = APIRouter(tags=["e1-ocr"])

LINE_SCHEMA = {
    "date": "交易日期 YYYY-MM-DD",
    "summary": "摘要/附言",
    "counterparty": "对方户名/收付款方",
    "amount": "金额（正数）",
    "direction": "收入 或 支出",
}

_EMPTY_FIELDS: dict[str, Any] = {
    "bank": "",
    "accountNo": "",
    "periodStart": "",
    "periodEnd": "",
    "lines": [],
}

_LLM_SYSTEM_PROMPT = (
    "你是注册会计师审计助手，专长于解读中国商业银行对账单或网银流水。\n"
    "请从OCR文本中提取账户信息与交易明细，严格返回JSON：\n"
    "{\n"
    '  "bank": "开户银行名称",\n'
    '  "accountNo": "银行账号",\n'
    '  "periodStart": "流水期间起 YYYY-MM-DD",\n'
    '  "periodEnd": "流水期间止 YYYY-MM-DD",\n'
    '  "lines": [ '
    + json.dumps(LINE_SCHEMA, ensure_ascii=False)
    + " ]\n"
    "}\n"
    "规则：\n"
    "1. lines 尽量提取全部可见交易；不要虚构OCR未出现的交易。\n"
    "2. amount 一律取绝对值正数；direction 用「收入」或「支出」。\n"
    "3. 贷方/转入/收款 → 收入；借方/转出/付款 → 支出。\n"
    "4. 日期统一 YYYY-MM-DD；无法判断留空。\n"
    "5. 若OCR过长导致无法全部提取，优先提取金额较大的交易，最多200笔。"
)


class E1StatementOcrResponse(BaseModel):
    attachment_id: str
    ocr_text: str
    extracted_fields: dict
    confidence: float
    file_name: str = ""
    line_count: int = 0
    skipped_line_count: int = 0
    reused: bool = False
    written_back: bool = False
    governed: bool = False
    requires_human_confirmation: bool = True


def _parse_llm_json(raw: str) -> dict:
    json_str = (raw or "").strip()
    if "```json" in json_str:
        json_str = json_str.split("```json", 1)[1].split("```", 1)[0].strip()
    elif "```" in json_str:
        json_str = json_str.split("```", 1)[1].split("```", 1)[0].strip()
    parsed = json.loads(json_str)
    return parsed if isinstance(parsed, dict) else {}


def _safe_float(v: Any) -> float:
    try:
        if v is None or v == "":
            return 0.0
        return abs(float(str(v).replace(",", "").replace("，", "").strip()))
    except (TypeError, ValueError):
        return 0.0


def _normalize_direction(raw: Any, amount_signed: float | None = None) -> str:
    s = str(raw or "").strip()
    if any(k in s for k in ("收", "入", "贷", "转入", "CR", "C", "credit")):
        return "收入"
    if any(k in s for k in ("支", "出", "借", "转出", "DR", "D", "debit")):
        return "支出"
    if amount_signed is not None and amount_signed < 0:
        return "支出"
    return ""


def _normalize_lines(raw: Any) -> tuple[list[dict[str, Any]], int]:
    if isinstance(raw, str):
        try:
            raw = json.loads(raw)
        except (json.JSONDecodeError, TypeError):
            return [], 0
    if not isinstance(raw, list):
        return [], 0
    rows: list[dict[str, Any]] = []
    skipped = 0
    overflow = max(0, len(raw) - 200)
    for item in raw[:200]:
        if not isinstance(item, dict):
            skipped += 1
            continue
        amount = _safe_float(item.get("amount"))
        direction = _normalize_direction(item.get("direction"), None)
        row = {
            "date": str(item.get("date") or "").strip(),
            "summary": str(item.get("summary") or "").strip(),
            "counterparty": str(item.get("counterparty") or "").strip(),
            "amount": amount,
            "direction": direction or ("收入" if amount else ""),
        }
        if not row["date"] and not row["amount"] and not row["counterparty"]:
            skipped += 1
            continue
        rows.append(row)
    return rows, skipped + overflow


def _normalize_fields(parsed: dict) -> tuple[dict[str, Any], int]:
    lines, skipped = _normalize_lines(parsed.get("lines"))
    return {
        "bank": str(parsed.get("bank") or "").strip(),
        "accountNo": str(parsed.get("accountNo") or "").strip(),
        "periodStart": str(parsed.get("periodStart") or "").strip(),
        "periodEnd": str(parsed.get("periodEnd") or "").strip(),
        "lines": lines,
        "skippedLineCount": skipped,
    }, skipped


@router.post("/api/workpapers/{wp_id}/e1/statement-ocr", response_model=E1StatementOcrResponse)
async def e1_statement_ocr(
    wp_id: str,
    file: UploadFile = File(...),
    attachment_id: str | None = Form(None),
    force_reocr: bool = Form(False),
    db: AsyncSession = Depends(get_db),
    _user: User = Depends(get_current_user),
) -> E1StatementOcrResponse:
    """上传银行对账单/流水，OCR + AI 提取明细供 E1-31 确认回填。"""
    try:
        wp_uuid = UUID(wp_id)
    except ValueError:
        raise HTTPException(400, "无效的 wp_id")

    linked_att_id = parse_optional_uuid(attachment_id)

    allowed = (".pdf", ".png", ".jpg", ".jpeg", ".tif", ".tiff", ".bmp", ".webp", ".xlsx", ".xls")
    filename = file.filename or "bank-statement.pdf"
    suffix = Path(filename).suffix.lower()
    if suffix not in allowed:
        raise HTTPException(400, f"不支持的文件类型: {suffix}，仅支持 {allowed}")

    if linked_att_id is not None and not force_reocr:
        reused = await load_reusable_ocr(db, linked_att_id)
        if reused is not None:
            fields = reused["extracted_fields"] or {}
            lines = fields.get("lines") if isinstance(fields, dict) else []
            skipped = fields.get("skippedLineCount", 0) if isinstance(fields, dict) else 0
            return E1StatementOcrResponse(
                attachment_id=str(linked_att_id),
                ocr_text=reused["ocr_text"],
                extracted_fields=fields,
                confidence=float(reused.get("confidence") or 0),
                file_name=filename,
                line_count=len(lines) if isinstance(lines, list) else 0,
                skipped_line_count=int(skipped or 0),
                reused=True,
                written_back=False,
            )

    content = await file.read()
    out_attachment_id = str(linked_att_id) if linked_att_id else str(uuid4())
    if linked_att_id is None:
        out_attachment_id, linked_att_id = await provision_wp_linked_attachment(
            db,
            wp_id=wp_uuid,
            file_name=filename,
            content=content,
            attachment_type="bank_statement",
            created_by=getattr(_user, "id", None),
        )

    storage_dir = Path("storage/workpapers") / wp_id / "bank-statements"
    storage_dir.mkdir(parents=True, exist_ok=True)
    file_path = storage_dir / f"{out_attachment_id}{suffix}"

    async with aiofiles.open(str(file_path), "wb") as f:
        await f.write(content)

    ocr_text = ""
    try:
        # Excel：直接尝试读单元格拼成文本，失败再走 OCR
        if suffix in (".xlsx", ".xls"):
            try:
                from openpyxl import load_workbook

                wb = load_workbook(str(file_path), data_only=True, read_only=True)
                chunks: list[str] = []
                for ws in wb.worksheets[:3]:
                    for row in ws.iter_rows(max_row=500, values_only=True):
                        cells = [str(c).strip() for c in row if c is not None and str(c).strip()]
                        if cells:
                            chunks.append("\t".join(cells))
                wb.close()
                ocr_text = "\n".join(chunks)
            except Exception as e:  # noqa: BLE001
                logger.warning("E1 statement excel parse failed: %s", e)
        if not ocr_text.strip():
            ocr_result = await UnifiedOCRService().recognize(str(file_path))
            ocr_text = ocr_result.get("text", "") or ""
    except Exception as e:  # noqa: BLE001
        logger.warning("E1 statement OCR failed for %s: %s", wp_id, e)
        if linked_att_id is not None:
            await AttachmentService(db).update_ocr_status(linked_att_id, "failed", ocr_text="")
            await db.commit()
        return E1StatementOcrResponse(
            attachment_id=out_attachment_id,
            ocr_text="",
            extracted_fields=dict(_EMPTY_FIELDS),
            confidence=0,
            file_name=filename,
            line_count=0,
            skipped_line_count=0,
        )

    if not ocr_text.strip():
        if linked_att_id is not None:
            await AttachmentService(db).update_ocr_status(linked_att_id, "failed", ocr_text="")
            await db.commit()
        return E1StatementOcrResponse(
            attachment_id=out_attachment_id,
            ocr_text="",
            extracted_fields=dict(_EMPTY_FIELDS),
            confidence=0,
            file_name=filename,
            line_count=0,
            skipped_line_count=0,
        )

    extracted_fields = dict(_EMPTY_FIELDS)
    confidence = 0.0
    skipped_line_count = 0
    try:
        messages = [
            {"role": "system", "content": _LLM_SYSTEM_PROMPT},
            {"role": "user", "content": f"OCR文本：\n{ocr_text[:16000]}"},
        ]
        llm_result = await chat_completion(messages=messages, temperature=0.1, max_tokens=8000)
        raw = llm_result if isinstance(llm_result, str) else str(llm_result)
        extracted_fields, skipped_line_count = _normalize_fields(_parse_llm_json(raw))
        lines = extracted_fields.get("lines") or []
        filled = (
            (1 if extracted_fields.get("bank") else 0)
            + (1 if extracted_fields.get("accountNo") else 0)
            + (1 if lines else 0)
        )
        confidence = round(min(1.0, filled / 3 + min(len(lines), 20) * 0.02), 2)
    except Exception as e:  # noqa: BLE001
        logger.warning("E1 statement LLM extraction failed for %s: %s", wp_id, e)
        if linked_att_id is not None:
            await AttachmentService(db).update_ocr_status(linked_att_id, "failed", ocr_text=ocr_text)
            await db.commit()
        return E1StatementOcrResponse(
            attachment_id=out_attachment_id,
            ocr_text=ocr_text,
            extracted_fields=dict(_EMPTY_FIELDS),
            confidence=0,
            file_name=filename,
            line_count=0,
            skipped_line_count=0,
        )

    lines = extracted_fields.get("lines") or []
    out_attachment_id, written = await finalize_linked_ocr_writeback(
        db,
        wp_id=wp_uuid,
        linked_att_id=linked_att_id,
        temp_id=out_attachment_id,
        ocr_text=ocr_text,
        extracted_fields=extracted_fields,
        confidence=confidence,
    )
    return E1StatementOcrResponse(
        attachment_id=out_attachment_id,
        ocr_text=ocr_text,
        extracted_fields=extracted_fields,
        confidence=confidence,
        file_name=filename,
        line_count=len(lines),
        skipped_line_count=skipped_line_count,
        written_back=written,
    )

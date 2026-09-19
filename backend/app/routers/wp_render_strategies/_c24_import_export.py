"""C24 会计分录细节测试 — 分录明细导入导出端点

3个端点：
- GET  /api/workpapers/{wp_id}/c24-journal/export-template  空白模板xlsx
- GET  /api/workpapers/{wp_id}/c24-journal/export-data      数据xlsx（已导入的分录）
- POST /api/workpapers/{wp_id}/c24-journal/import-data      解析xlsx写入分录明细

分录结构 JournalEntry（13列）：
凭证日期/凭证月份/凭证类型/凭证编号/摘要/科目编号/科目名称/借方金额/贷方金额/
凭证张数/填制人/审核人/记账人

数据存储：checklist_responses.remark（item_id = "C24-journal-entries"），JSON 数组。
"""

from __future__ import annotations

import io
import json
import logging
from typing import Any
from urllib.parse import quote
from uuid import uuid4

from fastapi import APIRouter, Depends, File, HTTPException, UploadFile
from fastapi.responses import StreamingResponse
from openpyxl import Workbook, load_workbook
from sqlalchemy.ext.asyncio import AsyncSession

from app.core.database import get_db
from app.deps import get_current_user
from app.models.core import User

logger = logging.getLogger(__name__)

router = APIRouter(tags=["c24-import-export"])

# ═══════════════════════════════════════════════════════════════════════════════
# 配置
# ═══════════════════════════════════════════════════════════════════════════════

_ROW_LIMIT = 20000  # 分录量可能很大（源模板虚拟数据 10541 行）

_ITEM_ID = "C24-journal-entries"

_HEADERS: list[str] = [
    "凭证日期",
    "凭证月份",
    "凭证类型",
    "凭证编号",
    "摘要",
    "科目编号",
    "科目名称",
    "借方金额",
    "贷方金额",
    "凭证张数",
    "填制人",
    "审核人",
    "记账人",
]

_HEADER_TO_FIELD: dict[str, str] = {
    "凭证日期": "voucherDate",
    "凭证月份": "voucherMonth",
    "凭证类型": "voucherType",
    "凭证编号": "voucherNo",
    "摘要": "summary",
    "科目编号": "accountCode",
    "科目名称": "accountName",
    "借方金额": "debit",
    "贷方金额": "credit",
    "凭证张数": "voucherSheets",
    "填制人": "preparer",
    "审核人": "reviewer",
    "记账人": "poster",
}

_FIELD_TO_HEADER: dict[str, str] = {v: k for k, v in _HEADER_TO_FIELD.items()}

# 编制说明文本
_GUIDANCE: list[str] = [
    "C24 会计分录导入模板 — 编制说明",
    "",
    "1. 本模板用于导入被审计单位会计分录明细数据",
    "2. 列说明：",
    "   - 凭证日期：格式 YYYY-MM-DD（如 2025-01-15）",
    "   - 凭证月份：数字 1~12",
    "   - 凭证类型：付/收/转 等",
    "   - 凭证编号：如 付-001、转-0023",
    "   - 摘要：凭证摘要文字",
    "   - 科目编号：如 6001、1001.01",
    "   - 科目名称：如 主营业务收入、银行存款-工商银行",
    "   - 借方金额：数字，无金额填 0",
    "   - 贷方金额：数字，无金额填 0",
    "   - 凭证张数：附件张数",
    "   - 填制人：编制凭证的人员姓名",
    "   - 审核人：审核凭证的人员姓名",
    "   - 记账人：记账（过账）的人员姓名",
    "",
    "3. 注意事项：",
    "   - 第一行为表头，请勿修改",
    "   - 数据从第2行开始，每行一条分录明细",
    "   - 同一凭证号可有多行（一借多贷/一贷多借/多借多贷）",
    "   - 金额列请填写数字，不要带千位分隔符或货币符号",
    "   - 最大支持 20000 行数据",
]


# ═══════════════════════════════════════════════════════════════════════════════
# 辅助函数
# ═══════════════════════════════════════════════════════════════════════════════


def _safe_float(val: Any) -> float:
    """安全转换为 float，失败返回 0.0。"""
    if val is None:
        return 0.0
    try:
        return float(val)
    except (ValueError, TypeError):
        return 0.0


def _safe_str(val: Any) -> str:
    """安全转换为 str。"""
    if val is None:
        return ""
    return str(val).strip()


def _safe_int(val: Any) -> int:
    """安全转换为 int，失败返回 0。"""
    if val is None:
        return 0
    try:
        return int(float(val))
    except (ValueError, TypeError):
        return 0


def _parse_row(row: tuple, actual_headers: list[str]) -> dict:
    """解析一行数据为 JournalEntry dict。"""
    entry: dict = {}
    for col_idx, header in enumerate(actual_headers):
        field = _HEADER_TO_FIELD.get(header)
        if field is None:
            continue
        val = row[col_idx] if col_idx < len(row) else None
        if field in ("debit", "credit"):
            entry[field] = _safe_float(val)
        elif field == "voucherMonth":
            entry[field] = _safe_int(val)
        else:
            entry[field] = _safe_str(val)
    return entry


def _export_row(data: dict) -> list:
    """将 JournalEntry dict 转为一行 xlsx 值（按 _HEADERS 顺序）。"""
    result = []
    for header in _HEADERS:
        field = _HEADER_TO_FIELD.get(header, "")
        val = data.get(field, "")
        result.append(val)
    return result


# ═══════════════════════════════════════════════════════════════════════════════
# 端点
# ═══════════════════════════════════════════════════════════════════════════════


@router.get("/api/workpapers/{wp_id}/c24-journal/export-template")
async def c24_export_template(
    wp_id: str,
    current_user: User = Depends(get_current_user),
) -> StreamingResponse:
    """导出空白分录模板xlsx（含表头+编制说明，无数据行）。"""
    wb = Workbook()
    ws = wb.active
    ws.title = "分录明细"
    ws.append(_HEADERS)
    ws.freeze_panes = "A2"
    # 设置列宽
    col_widths = [12, 8, 8, 12, 30, 12, 20, 14, 14, 8, 10, 10, 10]
    for col_idx, width in enumerate(col_widths, 1):
        col_letter = chr(64 + col_idx) if col_idx <= 26 else "A"
        ws.column_dimensions[col_letter].width = width

    # 编制说明 sheet
    ws_guide = wb.create_sheet("编制说明")
    ws_guide.append(["编制说明"])
    ws_guide.append([])
    for line in _GUIDANCE:
        ws_guide.append([line])
    ws_guide.column_dimensions["A"].width = 80

    buffer = io.BytesIO()
    wb.save(buffer)
    buffer.seek(0)

    filename = "C24_分录明细_模板.xlsx"
    encoded_filename = quote(filename)
    return StreamingResponse(
        buffer,
        media_type="application/vnd.openxmlformats-officedocument.spreadsheetml.sheet",
        headers={"Content-Disposition": f"attachment; filename*=UTF-8''{encoded_filename}"},
    )


@router.get("/api/workpapers/{wp_id}/c24-journal/export-data")
async def c24_export_data(
    wp_id: str,
    db: AsyncSession = Depends(get_db),
    current_user: User = Depends(get_current_user),
) -> StreamingResponse:
    """导出当前已导入的分录数据xlsx。"""
    import sqlalchemy as sa_mod

    result = await db.execute(
        sa_mod.text(
            "SELECT remark FROM checklist_responses "
            "WHERE wp_id = :wp_id AND item_id = :item_id LIMIT 1"
        ),
        {"wp_id": wp_id, "item_id": _ITEM_ID},
    )
    row = result.fetchone()
    rows_data: list[dict] = []
    if row and row.remark:
        try:
            rows_data = json.loads(row.remark)
        except (json.JSONDecodeError, TypeError):
            pass

    wb = Workbook()
    ws = wb.active
    ws.title = "分录明细"
    ws.append(_HEADERS)
    ws.freeze_panes = "A2"

    for data_row in rows_data:
        ws.append(_export_row(data_row))

    buffer = io.BytesIO()
    wb.save(buffer)
    buffer.seek(0)

    filename = "C24_分录明细_数据.xlsx"
    encoded_filename = quote(filename)
    return StreamingResponse(
        buffer,
        media_type="application/vnd.openxmlformats-officedocument.spreadsheetml.sheet",
        headers={"Content-Disposition": f"attachment; filename*=UTF-8''{encoded_filename}"},
    )


@router.post("/api/workpapers/{wp_id}/c24-journal/import-data")
async def c24_import_data(
    wp_id: str,
    file: UploadFile = File(...),
    db: AsyncSession = Depends(get_db),
    current_user: User = Depends(get_current_user),
) -> dict[str, Any]:
    """解析上传的分录xlsx，校验格式，写入 checklist_responses。"""
    import sqlalchemy as sa_mod

    if not file.filename or not file.filename.endswith(".xlsx"):
        raise HTTPException(400, "请上传 .xlsx 格式文件")

    content = await file.read()
    if len(content) > 20 * 1024 * 1024:
        raise HTTPException(400, "文件大小不能超过20MB")

    try:
        wb = load_workbook(io.BytesIO(content), read_only=True, data_only=True)
    except Exception:
        raise HTTPException(400, "无法解析xlsx文件，请确认文件格式正确")

    ws = wb.active
    if ws is None:
        raise HTTPException(400, "xlsx文件中无活动工作表")

    # 验证列头
    actual_headers = [
        str(cell.value).strip() if cell.value else ""
        for cell in next(ws.iter_rows(min_row=1, max_row=1))
    ]

    errors: list[str] = []
    # 必须列检查（允许多余列，但核心列不可缺）
    required_headers = ["凭证日期", "凭证编号", "科目编号", "借方金额", "贷方金额"]
    missing_cols = [h for h in required_headers if h not in actual_headers]
    if missing_cols:
        errors.append(f"缺少必需列: {', '.join(missing_cols)}")

    if errors:
        wb.close()
        raise HTTPException(400, detail=errors)

    # 解析数据行
    rows_data: list[dict] = []
    truncated = False
    row_count = 0
    row_errors: list[str] = []

    for row in ws.iter_rows(min_row=2, values_only=True):
        if all(v is None for v in row):
            continue
        row_count += 1
        if row_count > _ROW_LIMIT:
            truncated = True
            break
        entry = _parse_row(row, actual_headers)
        # 基本校验：凭证编号不能为空
        if not entry.get("voucherNo"):
            row_errors.append(f"第{row_count + 1}行: 凭证编号为空")
            if len(row_errors) >= 10:
                row_errors.append("...更多错误已省略")
                break
            continue
        rows_data.append(entry)

    wb.close()

    if row_errors and not rows_data:
        raise HTTPException(400, detail=row_errors)

    # 写入 checklist_responses
    remark_json = json.dumps(rows_data, ensure_ascii=False)

    await db.execute(
        sa_mod.text("""
            INSERT INTO checklist_responses (id, wp_id, item_id, remark, updated_at)
            VALUES (:id, :wp_id, :item_id, :remark, NOW())
            ON CONFLICT (wp_id, item_id)
            DO UPDATE SET remark = :remark, updated_at = NOW()
        """),
        {"id": str(uuid4()), "wp_id": wp_id, "item_id": _ITEM_ID, "remark": remark_json},
    )
    await db.commit()

    result_data: dict[str, Any] = {
        "ok": True,
        "imported_count": len(rows_data),
        "errors": row_errors,
    }
    if truncated:
        result_data["warning"] = f"数据行数超过{_ROW_LIMIT}行限制，已截断"
        result_data["truncated"] = True

    return result_data

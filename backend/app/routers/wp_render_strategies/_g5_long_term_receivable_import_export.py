"""G5 长期应收款 — 导入导出（9 张表：G5-1 + G5-2~7 / G5-11~12）."""

from __future__ import annotations

from typing import Any
from uuid import uuid4

from ._cycle_import_export_common import (
    ROW_LIMIT,
    build_workbook_template,
    create_cycle_import_export_router,
    export_row_by_keys,
    parse_row_by_headers,
    parse_upload_xlsx,
    safe_float,
    safe_str,
)

# ── G5-1 审定表（keyed remark store，对齐 G2-1）──
_G5_1_HEADERS = [
    "行键", "项目", "区段",
    "期初未审数", "期初账项调整", "期初重分类调整",
    "期末未审数", "期末账项调整", "期末重分类调整", "原因分析",
]
_G5_1_KEYS = [
    "rowKey", "label", "section",
    "openingUnadjusted", "openingAJE", "openingRJE",
    "closingUnadjusted", "closingAJE", "closingRJE", "reasonAnalysis",
]
_G5_1_KEEP_KEYS = [
    "openingUnadjusted", "openingAJE", "openingRJE",
    "closingUnadjusted", "closingAJE", "closingRJE", "reasonAnalysis",
]

_G5_1_LEAF_ROWS = [
    ("gross-individual", "其中：单项计提坏账准备的长期应收款", "余额"),
    ("gross-collective-business", "业务类型组合", "余额"),
    ("gross-collective-customer", "客户类型组合", "余额"),
    ("gross-one-year", "减：1年内到期的长期应收款", "余额"),
    ("provision-individual", "其中：单项计提坏账准备", "坏账准备"),
    ("provision-collective-business", "业务类型组合", "坏账准备"),
    ("provision-collective-customer", "客户类型组合", "坏账准备"),
    ("provision-one-year", "减：1年内到期对应坏账", "坏账准备"),
    ("tb-amount", "试算平衡表数", "试算"),
]


def _build_g5_1_prefill():
    rows = []
    for row_key, label, section in _G5_1_LEAF_ROWS:
        rows.append([row_key, label, section, 0, 0, 0, 0, 0, 0, ""])
    return rows


_G5_2_HEADERS = [
    "序号", "债务人", "业务类型", "合同编号", "开始日", "到期日", "1年内到期",
    "合同总额", "已收回", "期末余额", "借方发生", "贷方发生", "关联方", "未实现收益", "净额",
    "1年内", "1-2年", "2-3年", "3-4年", "4-5年", "5年以上", "账龄合计", "备注",
]
_G5_2_KEYS = [
    "seq", "debtorName", "businessType", "contractNo", "startDate", "maturityDate", "isWithinOneYear",
    "contractAmount", "recoveredAmount", "closingBalance", "debitOccurrence", "creditOccurrence",
    "isRelatedParty", "unrealizedIncome", "netAmount",
    "aging1Year", "aging1to2", "aging2to3", "aging3to4", "aging4to5", "aging5Plus", "agingTotal", "remark",
]

# 对齐前端 useG5BadDebtDetail 滚动态 leaf（ECL 测算在 G5-9/G5-10）
_G5_3_HEADERS = [
    "序号", "类别", "项目/债务人", "组合细分",
    "期初未审", "期初账项调整", "本期计提", "其他增加",
    "转回", "转销/核销", "其他减少", "期末账项调整", "差异原因",
]
_G5_3_KEYS = [
    "seq", "category", "item", "portfolioType",
    "openingUnadjusted", "openingAdjustment", "provisionIncrease", "otherIncrease",
    "reversal", "writeOff", "otherDecrease", "closingAdjustment", "reason",
]

_G5_4_HEADERS = [
    "序号", "类别", "分录类型", "日期", "摘要", "报表项目", "科目代码", "科目名称",
    "借方", "贷方", "索引", "编制人", "备注",
]
_G5_4_KEYS = [
    "seq", "category", "entryType", "date", "description", "reportItem", "accountCode", "accountName",
    "debitAmount", "creditAmount", "indexRef", "preparedBy", "remark",
]

_G5_5_HEADERS = [
    "项目名称", "期次", "承租人", "内含利率", "期初应收", "期初未实现", "本期收款", "企业账面收益", "备注",
]
_G5_5_KEYS = [
    "projectName", "periodNo", "lessee", "implicitRate", "openingReceivable", "openingUnrealized",
    "periodCollection", "companyBookIncome", "remark",
]

_G5_6_HEADERS = [
    "项目名称", "期次", "应收总额", "公允价值", "实际利率", "期初应收", "本期收款", "备注",
]
_G5_6_KEYS = [
    "projectName", "periodNo", "contractTotal", "fairValue", "effectiveRate", "openingReceivable", "periodCollection", "remark",
]

_G5_7_HEADERS = [
    "序号", "债务人", "保理商", "金额", "方式", "终止确认", "判断依据", "结论", "索引",
]
_G5_7_KEYS = [
    "seq", "debtor", "factor", "amount", "method", "derecognition", "basis", "conclusion", "indexRef",
]

_G5_11_HEADERS = [
    "区段", "序号", "债务人", "累计计提", "转回金额", "核销金额", "审批状态", "关联交易", "原因", "索引",
]
_G5_11_KEYS = [
    "section", "seq", "debtor", "accumulatedProvision", "reversalAmount", "writeoffAmount",
    "approvalStatus", "isRelatedParty", "reason", "indexRef",
]

_G5_12_HEADERS = [
    "区段", "序号", "日期", "凭证编号", "业务内容", "对方科目", "对方明细",
    "借方金额", "贷方金额", "支持性文件",
    "核对1完整", "核对2批准", "核对3账务", "核对4成本", "核对5对手",
    "索引号", "是否异常", "备注", "来源",
]
_G5_12_KEYS = [
    "section", "seq", "date", "voucherNo", "businessContent", "offsetAccount", "offsetSubAccount",
    "debitAmount", "creditAmount", "supportingDoc",
    "check1", "check2", "check3", "check4", "check5",
    "indexNo", "abnormal", "remark", "source",
]

_G5_12_SHEET_OCC = "本期发生额检查"
_G5_12_SHEET_POST = "期后处置新增检查"
_G5_12_SHEET_META = "样本与结论"
_G5_12_CHECK_LABELS = (
    "原始凭证内容完整",
    "有授权批准",
    "账务处理正确",
    "初始成本计算正确",
    "还款人与交易对手核对一致",
)


def _g5_12_empty_payload() -> dict[str, Any]:
    return {
        "criteria": {
            "populationDebitCount": 0,
            "populationDebitAmount": 0,
            "populationCreditCount": 0,
            "populationCreditAmount": 0,
            "specificSample": "",
            "samplingPopulationCount": 0,
            "samplingPopulationAmount": 0,
            "sampleSize": 0,
            "samplingMethod": "货币单元抽样",
            "samplingProcess": "",
            "bookDebitOccurrence": 0,
            "bookCreditOccurrence": 0,
        },
        "occurrenceRows": [],
        "postPeriodRows": [],
        "auditNote": "",
        "conclusion": "",
        "conclusionOption": "",
    }


def _g5_12_row_to_excel(row: dict, section: str, seq: int) -> list[Any]:
    checks = row.get("checks") or [False] * 5
    if not isinstance(checks, list):
        checks = [False] * 5
    while len(checks) < 5:
        checks.append(False)
    marks = ["✓" if c else "" for c in checks[:5]]
    abnormal = row.get("abnormal")
    if isinstance(abnormal, bool):
        abn = "是" if abnormal else "否"
    else:
        abn = str(abnormal or "否")
    return [
        section,
        seq,
        row.get("date") or row.get("voucherDate") or "",
        row.get("voucherNo") or "",
        row.get("businessContent") or row.get("summary") or "",
        row.get("offsetAccount") or row.get("counterAccount") or "",
        row.get("offsetSubAccount") or "",
        row.get("debitAmount") or 0,
        row.get("creditAmount") or 0,
        row.get("supportingDoc") or "",
        *marks,
        row.get("indexNo") or row.get("indexRef") or "",
        abn,
        row.get("remark") or row.get("conclusion") or "",
        row.get("source") or "",
    ]


def _g5_12_excel_to_row(vals: list[Any], seq: int) -> dict[str, Any]:
    from uuid import uuid4

    def g(i: int, default: Any = ""):
        return vals[i] if i < len(vals) and vals[i] is not None else default

    def yn(v: Any) -> bool:
        s = str(v or "").strip().lower()
        return s in {"是", "y", "yes", "true", "1", "异常"}

    def ck(v: Any) -> bool:
        s = str(v or "").strip()
        return s in {"✓", "√", "是", "Y", "y", "1", "true", "TRUE"}

    try:
        debit = float(g(7, 0) or 0)
    except (TypeError, ValueError):
        debit = 0.0
    try:
        credit = float(g(8, 0) or 0)
    except (TypeError, ValueError):
        credit = 0.0

    return {
        "id": str(uuid4()),
        "date": str(g(2)),
        "voucherNo": str(g(3)),
        "businessContent": str(g(4)),
        "offsetAccount": str(g(5)),
        "offsetSubAccount": str(g(6)),
        "debitAmount": debit,
        "creditAmount": credit,
        "supportingDoc": str(g(9)),
        "checks": [ck(g(10 + i)) for i in range(5)],
        "indexNo": str(g(15)),
        "abnormal": yn(g(16)),
        "remark": str(g(17)),
        "source": str(g(18) or "导入"),
    }


def _g5_12_write_header(ws) -> None:
    from openpyxl.styles import Font

    ws.append(_G5_12_HEADERS)
    for cell in ws[1]:
        cell.font = Font(bold=True)
    legend = [""] * len(_G5_12_HEADERS)
    legend[10] = "1." + _G5_12_CHECK_LABELS[0]
    legend[11] = "2." + _G5_12_CHECK_LABELS[1]
    legend[12] = "3." + _G5_12_CHECK_LABELS[2]
    legend[13] = "4." + _G5_12_CHECK_LABELS[3]
    legend[14] = "5." + _G5_12_CHECK_LABELS[4]
    ws.append(legend)


def _build_g5_12_workbook(payload: Any, template_only: bool = False):
    from openpyxl import Workbook
    from openpyxl.styles import Font

    data = payload if isinstance(payload, dict) else _g5_12_empty_payload()
    if isinstance(payload, list):
        data = {**_g5_12_empty_payload(), "occurrenceRows": payload}

    wb = Workbook()
    ws1 = wb.active
    ws1.title = _G5_12_SHEET_OCC
    _g5_12_write_header(ws1)
    if not template_only:
        for i, r in enumerate(data.get("occurrenceRows") or [], start=1):
            ws1.append(_g5_12_row_to_excel(r, "本期", i))

    ws2 = wb.create_sheet(_G5_12_SHEET_POST)
    _g5_12_write_header(ws2)
    if not template_only:
        for i, r in enumerate(data.get("postPeriodRows") or [], start=1):
            ws2.append(_g5_12_row_to_excel(r, "期后", i))

    ws3 = wb.create_sheet(_G5_12_SHEET_META)
    c = data.get("criteria") or {}
    meta_rows = [
        ("字段", "值"),
        ("测试总体借方笔数", c.get("populationDebitCount", 0)),
        ("测试总体借方金额", c.get("populationDebitAmount", 0)),
        ("测试总体贷方笔数", c.get("populationCreditCount", 0)),
        ("测试总体贷方金额", c.get("populationCreditAmount", 0)),
        ("特定样本", c.get("specificSample", "")),
        ("抽样总体笔数", c.get("samplingPopulationCount", 0)),
        ("抽样总体金额", c.get("samplingPopulationAmount", 0)),
        ("抽样样本量", c.get("sampleSize", 0)),
        ("抽样方法", c.get("samplingMethod", "")),
        ("抽样过程", c.get("samplingProcess", "")),
        ("账面本期借方", c.get("bookDebitOccurrence", 0)),
        ("账面本期贷方", c.get("bookCreditOccurrence", 0)),
        ("审计说明", data.get("auditNote", "")),
        ("结论选项", data.get("conclusionOption", "")),
        ("审计结论", data.get("conclusion", "")),
    ]
    for row in meta_rows:
        ws3.append(list(row))
    ws3["A1"].font = Font(bold=True)
    ws3["B1"].font = Font(bold=True)

    ws4 = wb.create_sheet("编制说明")
    for line in [
        "G5-12 长期应收款凭证检查表 — 导入导出说明",
        "",
        f"Sheet「{_G5_12_SHEET_OCC}」：本期发生额检查凭证明细",
        f"Sheet「{_G5_12_SHEET_POST}」：期后处置、新增检查凭证明细",
        f"Sheet「{_G5_12_SHEET_META}」：样本选取标准与审计说明/结论",
        "",
        "五项核对（填 ✓ 表示已核对通过）：",
        *[f"  {i + 1}. {lab}" for i, lab in enumerate(_G5_12_CHECK_LABELS)],
        "",
        "是否异常填：是/否",
        "导入保留样本与结论区并替换两区明细；旧版单表扁平导入兼容（全部归入本期）。",
        "检查比例 = 检查金额 ÷ 账面金额（账面来自「账面本期借/贷方」）。",
    ]:
        ws4.append([line])

    return wb


def _parse_g5_12_import(content: bytes) -> tuple[dict[str, Any], list[str], int]:
    from io import BytesIO

    from openpyxl import load_workbook

    errors: list[str] = []
    wb = load_workbook(BytesIO(content), data_only=True)
    payload = _g5_12_empty_payload()
    count = 0

    def read_detail(sheet_name: str) -> list[dict]:
        nonlocal count
        if sheet_name not in wb.sheetnames:
            return []
        ws = wb[sheet_name]
        rows_out: list[dict] = []
        for i, row in enumerate(ws.iter_rows(min_row=3, values_only=True), start=1):
            vals = list(row or [])
            if not any(v is not None and str(v).strip() != "" for v in vals[2:9]):
                continue
            rows_out.append(_g5_12_excel_to_row(vals, i))
            count += 1
            if count >= 500:
                errors.append("超过500行限制，已截断")
                break
        return rows_out

    if _G5_12_SHEET_OCC in wb.sheetnames or _G5_12_SHEET_POST in wb.sheetnames:
        payload["occurrenceRows"] = read_detail(_G5_12_SHEET_OCC)
        payload["postPeriodRows"] = read_detail(_G5_12_SHEET_POST)
    else:
        first = wb.sheetnames[0]
        payload["occurrenceRows"] = read_detail(first)

    if _G5_12_SHEET_META in wb.sheetnames:
        ws = wb[_G5_12_SHEET_META]
        cmap: dict[str, Any] = {}
        for row in ws.iter_rows(min_row=2, values_only=True):
            if not row or row[0] is None:
                continue
            cmap[str(row[0]).strip()] = row[1] if len(row) > 1 else ""
        crit = payload["criteria"]
        mapping = {
            "测试总体借方笔数": "populationDebitCount",
            "测试总体借方金额": "populationDebitAmount",
            "测试总体贷方笔数": "populationCreditCount",
            "测试总体贷方金额": "populationCreditAmount",
            "特定样本": "specificSample",
            "抽样总体笔数": "samplingPopulationCount",
            "抽样总体金额": "samplingPopulationAmount",
            "抽样样本量": "sampleSize",
            "抽样方法": "samplingMethod",
            "抽样过程": "samplingProcess",
            "账面本期借方": "bookDebitOccurrence",
            "账面本期贷方": "bookCreditOccurrence",
        }
        for zh, key in mapping.items():
            if zh in cmap and cmap[zh] is not None:
                if key in {"specificSample", "samplingMethod", "samplingProcess"}:
                    crit[key] = str(cmap[zh])
                else:
                    try:
                        crit[key] = float(cmap[zh] or 0)
                    except (TypeError, ValueError):
                        crit[key] = 0
        if "审计说明" in cmap:
            payload["auditNote"] = str(cmap["审计说明"] or "")
        if "审计结论" in cmap:
            payload["conclusion"] = str(cmap["审计结论"] or "")
        if "结论选项" in cmap:
            payload["conclusionOption"] = str(cmap["结论选项"] or "")

    if count == 0 and not payload["auditNote"] and not payload["conclusion"]:
        errors.append("未解析到有效凭证行或样本结论")
    return payload, errors, count


def _rate_to_percent(value: Any) -> float:
    num = safe_float(value)
    if abs(num) <= 1:
        return round(num * 100, 6)
    return num


def _rate_to_decimal(value: Any) -> float:
    num = safe_float(value)
    if abs(num) > 1:
        return round(num / 100, 8)
    return num


def _flatten_lease_like_groups(payload: Any, *, rate_key: str) -> list[dict]:
    """嵌套 {groups:[{projectName, basic, periods}]} → 扁平行；利率输出百分数。"""
    if isinstance(payload, list):
        # 已是扁平行或 group 列表
        if payload and isinstance(payload[0], dict) and "periods" not in payload[0] and "basic" not in payload[0]:
            flat = []
            for row in payload:
                r = dict(row)
                if rate_key in r:
                    r[rate_key] = _rate_to_percent(r.get(rate_key))
                flat.append(r)
            return flat
        groups = payload
    elif isinstance(payload, dict):
        groups = payload.get("groups") if isinstance(payload.get("groups"), list) else []
    else:
        return []
    flat: list[dict] = []
    for group in groups:
        if not isinstance(group, dict):
            continue
        basic = group.get("basic") if isinstance(group.get("basic"), dict) else {}
        periods = group.get("periods") if isinstance(group.get("periods"), list) else [{}]
        if not periods:
            periods = [{}]
        for idx, period in enumerate(periods):
            period = period if isinstance(period, dict) else {}
            row = {
                "projectName": group.get("projectName") or "",
                "periodNo": period.get("periodNo") or idx + 1,
                "lessee": basic.get("lessee") or group.get("lessee") or "",
                "implicitRate": _rate_to_percent(basic.get("implicitRate", group.get("implicitRate"))),
                "effectiveRate": _rate_to_percent(basic.get("effectiveRate", group.get("effectiveRate"))),
                "contractTotal": basic.get("contractTotal", group.get("contractTotal")),
                "fairValue": basic.get("fairValue", group.get("fairValue")),
                "openingReceivable": period.get("openingReceivable"),
                "openingUnrealized": period.get("openingUnrealized"),
                "periodCollection": period.get("periodCollection"),
                "companyBookIncome": period.get("companyBookIncome"),
                "remark": period.get("remark") or group.get("remark") or "",
                "id": group.get("id") if idx == 0 else f"{group.get('id')}-{period.get('id') or idx}",
            }
            flat.append(row)
    return flat


def _nest_lease_like_rows(rows: list[dict], *, rate_key: str) -> dict[str, Any]:
    """扁平行 → {groups:[...] }；利率存小数。"""
    grouped: dict[str, dict] = {}
    order: list[str] = []
    for row in rows or []:
        name = safe_str(row.get("projectName")) or "未命名项目"
        key = "".join(name.split())
        if key not in grouped:
            basic = {
                "lessee": safe_str(row.get("lessee")),
                "fairValue": safe_float(row.get("fairValue")),
                "contractTotal": safe_float(row.get("contractTotal")),
                rate_key: _rate_to_decimal(row.get(rate_key)),
            }
            grouped[key] = {
                "id": safe_str(row.get("id")) or str(uuid4()),
                "projectName": name,
                "basic": basic,
                "periods": [],
            }
            order.append(key)
        grouped[key]["periods"].append({
            "id": str(uuid4()),
            "periodNo": int(safe_float(row.get("periodNo")) or len(grouped[key]["periods"]) + 1),
            "openingReceivable": safe_float(row.get("openingReceivable")),
            "openingUnrealized": safe_float(row.get("openingUnrealized")),
            "periodCollection": safe_float(row.get("periodCollection")),
            "companyBookIncome": safe_float(row.get("companyBookIncome")),
            "remark": safe_str(row.get("remark")),
        })
    return {"groups": [grouped[k] for k in order], "conclusion": ""}


def _build_g5_5_workbook(payload: Any, *, template_only: bool = False):
    rows = [] if template_only else _flatten_lease_like_groups(payload, rate_key="implicitRate")
    wb = build_workbook_template(
        "G5-5",
        _G5_5_HEADERS,
        title="G5-5 融资租赁测算",
        guidance=[
            "G5-5 融资租赁",
            "",
            "内含利率填百分数（如 5.25 表示 5.25%）；系统存储为小数。",
            "同一项目名称多行表示多期；导入后归并为 groups。",
        ],
    )
    ws = wb["G5-5"]
    for d in rows:
        ws.append(export_row_by_keys(d, _G5_5_KEYS))
    return wb


def _parse_g5_5_import(content: bytes):
    actual, raw = parse_upload_xlsx(content, _G5_5_HEADERS, header_row=2)
    rows: list[dict] = []
    errors: list[str] = []
    for i, r in enumerate(raw, start=1):
        if i > ROW_LIMIT:
            errors.append(f"数据行超过{ROW_LIMIT}行限制，已截断")
            break
        rows.append(parse_row_by_headers(r, actual, _G5_5_KEYS))
    payload = _nest_lease_like_rows(rows, rate_key="implicitRate")
    return payload, errors, len(rows)


def _build_g5_6_workbook(payload: Any, *, template_only: bool = False):
    rows = [] if template_only else _flatten_lease_like_groups(payload, rate_key="effectiveRate")
    wb = build_workbook_template(
        "G5-6",
        _G5_6_HEADERS,
        title="G5-6 分期销售测算",
        guidance=[
            "G5-6 分期销售",
            "",
            "实际利率填百分数（如 5.25 表示 5.25%）；系统存储为小数。",
            "同一项目名称多行表示多期；导入后归并为 groups。",
        ],
    )
    ws = wb["G5-6"]
    for d in rows:
        ws.append(export_row_by_keys(d, _G5_6_KEYS))
    return wb


def _parse_g5_6_import(content: bytes):
    actual, raw = parse_upload_xlsx(content, _G5_6_HEADERS, header_row=2)
    rows: list[dict] = []
    errors: list[str] = []
    for i, r in enumerate(raw, start=1):
        if i > ROW_LIMIT:
            errors.append(f"数据行超过{ROW_LIMIT}行限制，已截断")
            break
        rows.append(parse_row_by_headers(r, actual, _G5_6_KEYS))
    payload = _nest_lease_like_rows(rows, rate_key="effectiveRate")
    return payload, errors, len(rows)


def _flatten_g5_11(payload: Any) -> list[dict]:
    if isinstance(payload, list):
        return [dict(r) for r in payload if isinstance(r, dict)]
    if not isinstance(payload, dict):
        return []
    flat: list[dict] = []
    for section, key in (("reversal", "reversal"), ("writeoff", "writeoff")):
        rows = payload.get(key) if isinstance(payload.get(key), list) else []
        for row in rows:
            if not isinstance(row, dict):
                continue
            flat.append({
                "section": section,
                "seq": row.get("seq"),
                "debtor": row.get("debtor") or row.get("unitName") or "",
                "accumulatedProvision": row.get("accumulatedProvision"),
                "reversalAmount": row.get("reversalAmount") if section == "reversal" else "",
                "writeoffAmount": row.get("writeoffAmount") if section == "writeoff" else row.get("writeOffAmount"),
                "approvalStatus": row.get("approvalStatus") or ("是" if row.get("approvalComplete") else "否"),
                "isRelatedParty": "是" if row.get("isRelatedParty") else "否",
                "reason": row.get("reason") or row.get("reasonAnalysis") or "",
                "indexRef": row.get("indexRef") or "",
            })
    return flat


def _nest_g5_11_rows(rows: list[dict]) -> dict[str, Any]:
    reversal: list[dict] = []
    writeoff: list[dict] = []
    for row in rows or []:
        section = safe_str(row.get("section")).lower()
        related = safe_str(row.get("isRelatedParty")) in ("是", "true", "1", "yes")
        base = {
            "id": str(uuid4()),
            "seq": int(safe_float(row.get("seq")) or 0),
            "debtor": safe_str(row.get("debtor")),
            "accumulatedProvision": safe_float(row.get("accumulatedProvision")),
            "isRelatedParty": related,
            "reason": safe_str(row.get("reason")),
            "indexRef": safe_str(row.get("indexRef")),
        }
        if section in ("writeoff", "writeoffs", "核销"):
            writeoff.append({
                **base,
                "writeoffAmount": safe_float(row.get("writeoffAmount")),
                "approvalComplete": safe_str(row.get("approvalStatus")) in ("是", "true", "1", "yes", "已审批", "完成"),
            })
        else:
            reversal.append({
                **base,
                "reversalAmount": safe_float(row.get("reversalAmount")),
            })
    return {"reversal": reversal, "writeoff": writeoff}


def _build_g5_11_workbook(payload: Any, *, template_only: bool = False):
    rows = [] if template_only else _flatten_g5_11(payload)
    wb = build_workbook_template(
        "G5-11",
        _G5_11_HEADERS,
        title="G5-11 转回核销检查",
        guidance=["G5-11 转回核销", "", "区段填 reversal 或 writeoff；转回金额≤累计计提。"],
    )
    ws = wb["G5-11"]
    for d in rows:
        ws.append(export_row_by_keys(d, _G5_11_KEYS))
    return wb


def _parse_g5_11_import(content: bytes):
    actual, raw = parse_upload_xlsx(content, _G5_11_HEADERS, header_row=2)
    rows: list[dict] = []
    errors: list[str] = []
    for i, r in enumerate(raw, start=1):
        if i > ROW_LIMIT:
            errors.append(f"数据行超过{ROW_LIMIT}行限制，已截断")
            break
        rows.append(parse_row_by_headers(r, actual, _G5_11_KEYS))
    return _nest_g5_11_rows(rows), errors, len(rows)


_G5_SPECS: dict[str, dict[str, Any]] = {
    "G5-1": {
        "item_id": "G5-1-rows",
        "title": "G5-1 长期应收款审定表",
        "headers": _G5_1_HEADERS,
        "field_keys": _G5_1_KEYS,
        "storage_field": "conclusion",
        "dual_write": True,
        "keyed_by": "rowKey",
        "keep_keys": _G5_1_KEEP_KEYS,
        "template_prefill": _build_g5_1_prefill(),
        "guidance": [
            "G5-1 审定表 编制说明",
            "",
            "行键(rowKey)须与模板一致，勿改动；仅填写未审/账项调整/重分类调整/原因分析。",
            "区段：余额(gross) / 坏账准备(provision) / 试算；净额与一年以上列示数由前端自动计算。",
            "审定=未审+账项调整+重分类调整；|变动率|>30% 时原因分析必填。",
        ],
    },
    "G5-2": {
        "item_id": "G5-2-rows",
        "title": "G5-2 长期应收款余额明细表",
        "headers": _G5_2_HEADERS,
        "field_keys": _G5_2_KEYS,
        "storage_field": "conclusion",
        "dual_write": True,
        "aging": {
            "subject": "G5",
            "base_headers": [
                "序号", "债务人", "业务类型", "合同编号", "开始日", "到期日", "1年内到期",
                "合同总额", "已收回", "期末余额", "借方发生", "贷方发生", "关联方", "未实现收益", "净额", "备注",
            ],
            "base_field_keys": [
                "seq", "debtorName", "businessType", "contractNo", "startDate", "maturityDate", "isWithinOneYear",
                "contractAmount", "recoveredAmount", "closingBalance", "debitOccurrence", "creditOccurrence",
                "isRelatedParty", "unrealizedIncome", "netAmount", "remark",
            ],
        },
        "guidance": [
            "G5-2 余额明细",
            "",
            "期末余额=合同总额-已收回；净额=余额-未实现；账龄合计应等于净额。",
            "借方发生/贷方发生对齐源模板 H/I，供 G5-12 检查比例分母勾稽。",
            "1年内到期填 是/否 或 TRUE/FALSE，供 G5-1 汇总「减：1年内到期」。",
        ],
    },
    "G5-3": {
        "item_id": "G5-3-rows",
        "title": "G5-3 坏账准备明细表",
        "headers": _G5_3_HEADERS,
        "field_keys": _G5_3_KEYS,
        "storage_field": "conclusion",
        "dual_write": True,
        "guidance": [
            "G5-3 坏账准备明细（滚动）",
            "",
            "类别填 individual（单项）或 portfolio（组合）；组合细分填 business/customer。",
            "期初审定=期初未审+期初账项调整；期末未审=期初审定+计提+其他增加−转回−转销−其他减少；期末审定=期末未审+期末账项调整。",
            "ECL 测算请用 G5-9/G5-10，勿在本表填损失率。",
        ],
    },
    "G5-4": {
        "item_id": "G5-4-rows",
        "title": "G5-4 调整分录汇总",
        "headers": _G5_4_HEADERS,
        "field_keys": _G5_4_KEYS,
        "storage_field": "conclusion",
        "dual_write": True,
        "guidance": [
            "G5-4 调整分录",
            "",
            "类别填 账项调整/报表调整/其他；分录类型填 AJE 或 RJE（与类别联动）。",
            "摘要对应 description；借贷须平衡。",
        ],
    },
    "G5-5": {
        "item_id": "G5-5-rows",
        "title": "G5-5 融资租赁测算",
        "headers": _G5_5_HEADERS,
        "field_keys": _G5_5_KEYS,
        "storage_field": "conclusion",
        "dual_write": True,
        "build_workbook": _build_g5_5_workbook,
        "parse_import": _parse_g5_5_import,
        "guidance": [
            "G5-5 融资租赁",
            "",
            "内含利率填百分数（如 5.25 表示 5.25%）。",
        ],
    },
    "G5-6": {
        "item_id": "G5-6-rows",
        "title": "G5-6 分期销售测算",
        "headers": _G5_6_HEADERS,
        "field_keys": _G5_6_KEYS,
        "storage_field": "conclusion",
        "dual_write": True,
        "build_workbook": _build_g5_6_workbook,
        "parse_import": _parse_g5_6_import,
        "guidance": [
            "G5-6 分期销售",
            "",
            "实际利率填百分数（如 5.25 表示 5.25%）。",
        ],
    },
    "G5-7": {
        "item_id": "G5-7-rows",
        "title": "G5-7 保理核查表",
        "headers": _G5_7_HEADERS,
        "field_keys": _G5_7_KEYS,
        "storage_field": "conclusion",
        "dual_write": True,
        "guidance": ["G5-7 保理核查", "", "有追索权保理通常不应终止确认。"],
    },
    "G5-11": {
        "item_id": "G5-11-rows",
        "title": "G5-11 转回核销检查",
        "headers": _G5_11_HEADERS,
        "field_keys": _G5_11_KEYS,
        "storage_field": "conclusion",
        "dual_write": True,
        "build_workbook": _build_g5_11_workbook,
        "parse_import": _parse_g5_11_import,
        "guidance": ["G5-11 转回核销", "", "区段填 reversal 或 writeoff；转回金额≤累计计提。"],
    },
    "G5-12": {
        "item_id": "G5-12-voucher-check",
        "title": "G5-12 凭证检查表",
        "headers": _G5_12_HEADERS,
        "field_keys": _G5_12_KEYS,
        "storage_field": "conclusion",
        "dual_write": True,
        "build_workbook": _build_g5_12_workbook,
        "parse_import": _parse_g5_12_import,
        "guidance": [
            "G5-12 凭证检查（双区：本期发生额 / 期后处置新增）",
            "",
            "五项核对填 ✓；是否异常填 是/否；样本与结论见独立 sheet。",
        ],
    },
}

router = create_cycle_import_export_router(
    tag="g5-import-export",
    api_prefix="g5",
    specs=_G5_SPECS,
    storage_field="conclusion",
)

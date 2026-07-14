"""J3 股份支付 — 导入导出端点.

三端点：template / export / import（参照 D4-2 范式：导出模板/导出数据/导入数据）。
J3-1 股份支付情况表为动态行表格，需要导入导出支持。

对齐源模板「股份支付情况表J3-1」列结构；持久化到 checklist_responses.remark
（item_id J3-1-plans，JSON 打包），与前端 J3TabDetail 一致。

Spec: .kiro/specs/j3-share-based-payment/
Requirements: 2.7
"""
from __future__ import annotations
import io
import json
import logging
from datetime import datetime
from urllib.parse import quote

from fastapi import APIRouter, Depends, UploadFile, File, HTTPException
from fastapi.responses import StreamingResponse
import sqlalchemy as sa

from app.core.database import get_db
from app.deps import get_current_user

logger = logging.getLogger(__name__)

router = APIRouter(prefix="/api/workpapers/{wp_id}/import-export", tags=["j3-import-export"])

# ── J3-1 情况表列定义（对齐源模板）：(表头, 字段名, 是否数值) ──────────────────────
# 序号列由导出时自动写入，导入时忽略。
J3_DETAIL_COLUMNS: list[tuple[str, str, bool]] = [
    ("股份支付项目名称", "name", False),
    ("类型", "type", False),
    ("授予日", "grantDate", False),
    ("批准部门", "approvalDept", False),
    ("行权日", "exerciseDate", False),
    ("权益工具数量", "instrumentQty", True),
    ("等待期", "vestingPeriod", False),
    ("公允价值确定方法和数据来源", "fvMethod", False),
    ("协议变更/取消情况", "agreementChange", False),
    ("资产负债表日估计更新情况", "bsUpdate", False),
    ("剩余等待期限", "remainingPeriod", False),
    ("协议索引号", "agreementIndex", False),
    ("股份支付计算表索引号", "calcTableIndex", False),
    ("结论", "conclusion", False),
]
_ITEM_ID = "J3-1-plans"


def _headers() -> list[str]:
    return ["序号"] + [c[0] for c in J3_DETAIL_COLUMNS]


def _stream_xlsx(wb, filename: str) -> StreamingResponse:
    buf = io.BytesIO()
    wb.save(buf)
    buf.seek(0)
    encoded = quote(filename)
    return StreamingResponse(
        buf,
        media_type="application/vnd.openxmlformats-officedocument.spreadsheetml.sheet",
        headers={"Content-Disposition": f"attachment; filename*=UTF-8''{encoded}"},
    )


@router.get("/template")
async def export_template(wp_id: str, user=Depends(get_current_user)):
    """导出空白模板（含表头 + 示例行）."""
    try:
        import openpyxl
        wb = openpyxl.Workbook()
        ws = wb.active
        ws.title = "股份支付情况表J3-1"
        for col_idx, header in enumerate(_headers(), 1):
            ws.cell(row=1, column=col_idx, value=header)
        # 示例行
        example = ["1", "（示例）2024年股票期权激励计划", "以权益工具结算", "2024-01-15",
                   "董事会/股东大会", "", 1000000, "3 年", "Black-Scholes 期权定价模型",
                   "无", "按最佳估计更新可行权数量", "2 年", "S12", "J3-2", "条款与工具一致，未见异常"]
        for col_idx, val in enumerate(example, 1):
            ws.cell(row=2, column=col_idx, value=val)
        return _stream_xlsx(wb, f"J3-1_股份支付情况表_模板_{datetime.now().strftime('%Y%m%d')}.xlsx")
    except Exception as e:
        logger.error("J3 template export failed: %s", e)
        raise HTTPException(status_code=500, detail="模板导出失败") from e


@router.get("/export")
async def export_data(wp_id: str, db=Depends(get_db), user=Depends(get_current_user)):
    """导出当前数据."""
    try:
        import openpyxl
        result = await db.execute(
            sa.text(
                "SELECT remark FROM checklist_responses "
                "WHERE wp_id = :wp_id AND item_id = :item_id"
            ),
            {"wp_id": wp_id, "item_id": _ITEM_ID},
        )
        row = result.fetchone()
        plans: list[dict] = []
        if row and row.remark:
            try:
                parsed = json.loads(row.remark)
                if isinstance(parsed, list):
                    plans = parsed
            except (json.JSONDecodeError, TypeError):
                pass

        wb = openpyxl.Workbook()
        ws = wb.active
        ws.title = "股份支付情况表J3-1"
        for col_idx, header in enumerate(_headers(), 1):
            ws.cell(row=1, column=col_idx, value=header)
        for row_idx, plan in enumerate(plans, 2):
            ws.cell(row=row_idx, column=1, value=row_idx - 1)  # 序号
            for offset, (_, field, _is_num) in enumerate(J3_DETAIL_COLUMNS, 2):
                ws.cell(row=row_idx, column=offset, value=plan.get(field, ""))
        return _stream_xlsx(wb, f"J3-1_股份支付情况表_数据_{datetime.now().strftime('%Y%m%d')}.xlsx")
    except Exception as e:
        logger.error("J3 export failed: %s", e)
        raise HTTPException(status_code=500, detail="数据导出失败") from e


@router.post("/import")
async def import_data(
    wp_id: str,
    file: UploadFile = File(...),
    db=Depends(get_db),
    user=Depends(get_current_user),
):
    """导入 Excel 数据（按位置解析，回写 J3-1-plans）."""
    try:
        import openpyxl
        content = await file.read()
        wb = openpyxl.load_workbook(io.BytesIO(content), data_only=True)
        ws = wb.active

        # 表头校验（第 2 列应为"股份支付项目名称"）
        header_cell = str(ws.cell(row=1, column=2).value or "").strip()
        if header_cell and "股份支付项目名称" not in header_cell and "项目名称" not in header_cell:
            raise HTTPException(status_code=400, detail="模板列名不匹配，请使用导出的 J3-1 模板")

        plans: list[dict] = []
        for row_idx in range(2, ws.max_row + 1):
            name = ws.cell(row=row_idx, column=2).value  # 第 2 列 = 项目名称（第 1 列为序号）
            if not name or not str(name).strip():
                continue
            plan: dict = {"id": row_idx - 1}
            for offset, (_, field, is_num) in enumerate(J3_DETAIL_COLUMNS, 2):
                raw = ws.cell(row=row_idx, column=offset).value
                if is_num:
                    try:
                        plan[field] = float(raw) if raw not in (None, "") else 0
                    except (ValueError, TypeError):
                        plan[field] = 0
                else:
                    plan[field] = str(raw).strip() if raw is not None else ""
            if not plan.get("type"):
                plan["type"] = "以权益工具结算"
            plans.append(plan)

        # 取 project_id（checklist_responses.project_id NOT NULL）
        pid_res = await db.execute(
            sa.text("SELECT project_id FROM working_paper WHERE id = :wp_id"),
            {"wp_id": wp_id},
        )
        pid_row = pid_res.fetchone()
        if not pid_row:
            raise HTTPException(status_code=404, detail="底稿不存在")
        project_id = pid_row.project_id

        await db.execute(
            sa.text(
                "INSERT INTO checklist_responses (project_id, wp_id, item_id, remark, conclusion) "
                "VALUES (:project_id, :wp_id, :item_id, :remark, NULL) "
                "ON CONFLICT (wp_id, item_id) DO UPDATE SET remark = :remark, updated_at = now()"
            ),
            {
                "project_id": project_id,
                "wp_id": wp_id,
                "item_id": _ITEM_ID,
                "remark": json.dumps(plans, ensure_ascii=False),
            },
        )
        await db.commit()
        return {"code": 0, "message": f"成功导入 {len(plans)} 个方案", "data": {"count": len(plans)}}
    except HTTPException:
        raise
    except Exception as e:
        logger.error("J3 import failed: %s", e)
        raise HTTPException(status_code=500, detail="数据导入失败") from e

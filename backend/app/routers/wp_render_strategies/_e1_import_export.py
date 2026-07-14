"""E1 货币资金动态行表的 Excel 导入导出端点。

字段和 checklist_responses item_id 以 E1 前端持久化 JSON 为 canonical；
历史后端字段仅用于导出兼容。计算字段随数据导出，但导入时始终忽略。
"""
from __future__ import annotations

import io
import json
import logging
from datetime import date, datetime
from typing import Any
from urllib.parse import quote
from uuid import uuid4

import sqlalchemy as sa
from fastapi import APIRouter, Depends, File, HTTPException, Query, UploadFile
from fastapi.responses import StreamingResponse
from openpyxl import Workbook, load_workbook
from openpyxl.styles import Alignment, Font, PatternFill
from openpyxl.utils.datetime import from_excel
from sqlalchemy.ext.asyncio import AsyncSession

from app.core.database import get_db
from app.deps import get_current_user
from app.models.core import User

logger = logging.getLogger(__name__)
router = APIRouter(tags=["e1-import-export"])

_ROW_LIMIT = 500
_ACCOUNT_CODE_WHITELIST: set[str] = {"1001", "1002", "1012"}

# 中文列名 -> 前端持久化 JSON canonical 字段。dict 顺序即 Excel 列顺序。
_FIELD_MAPS: dict[str, dict[str, str]] = {
    "E1-2": {
        "币种": "currency", "期初余额": "opening", "本期增加": "increase",
        "本期减少": "decrease", "期末原币": "endingFc", "期末汇率": "fxRate",
        "期末折算人民币": "endingRmb", "账项调整（原币）": "adjustment",
        "期末审定人民币": "auditedRmb", "备注": "note",
    },
    "E1-3": {
        "一级区段": "section", "二级分组": "group", "开户银行": "bankName",
        "总账银行名称": "totalLedgerBank", "账号": "accountNo",
        "账户性质/主要用途": "accountType", "期初余额": "opening",
        "本期增加": "increase", "本期减少": "decrease", "期末余额": "ending",
        "账项调整": "adjustment", "审定数": "audited", "对账单余额": "statementBalance",
        "账面与对账单差异": "accountStatementDiff", "银行对账单索引号": "statementIndexNo",
        "银行存款余额调节表索引号": "reconciliationIndexNo", "回函确认金额": "confirmAmount",
        "询证函索引号": "confirmIndexNo", "函证差异": "confirmDiff",
        "受限金额": "restrictedAmount", "受限原因": "restrictedReason", "利率": "interestRate",
        "备注": "note", "原币币种": "fxCurrency", "期末汇率": "fxRate",
        "期初原币": "openingFc", "增加原币": "increaseFc", "减少原币": "decreaseFc",
        "期末原币": "endingFc", "调整原币": "adjustmentFc", "审定原币": "auditedFc",
    },
    "E1-4": {
        "序号": "seq", "开户银行": "bankName", "币种": "currency", "汇率": "fxRate",
        "期初余额": "opening", "本期增加": "increase", "本期减少": "decrease",
        "期末原币": "endingFc", "期末本位币": "endingRmb", "账项调整": "adjustment",
        "审定原币": "auditedFc", "审定人民币": "auditedRmb", "查询余额": "queryBalance",
        "差异": "diff", "差异原因": "diffReason", "银行余额索引号": "indexNo",
        "银行询证函索引号": "confirmationIndexNo", "备注": "note",
    },
    "E1-5": {
        "调整事项说明": "description", "类别": "category", "报表项目": "reportItem",
        "科目名称": "accountName", "附注项目": "noteItem", "借方金额": "debit",
        "贷方金额": "credit", "索引号": "indexNo", "备注": "note",
    },
    "E1-6": {
        "开户银行": "bankName", "账号": "accountNo", "企业账面余额": "bookBalance",
        "银行对账单余额": "statementBalance",
        "银行已收企业未收明细(JSON)": "bankReceivedCompanyUnreceivedItems",
        "银行已付企业未付明细(JSON)": "bankPaidCompanyUnpaidItems",
        "企业已收银行未收明细(JSON)": "companyReceivedBankUnreceivedItems",
        "企业已付银行未付明细(JSON)": "companyPaidBankUnpaidItems",
        "企业已收银行未收合计": "bankReceived", "企业已付银行未付合计": "bankPaid",
        "银行已收企业未收合计": "companyReceived", "银行已付企业未付合计": "companyPaid",
        "调节后企业余额": "reconciledBook", "调节后银行余额": "reconciledStatement",
        "差异": "diff", "差异原因": "diffReason",
    },
    "E1-7": {"面值": "denomination", "张数": "quantity", "金额": "subtotal"},
    "E1-8": {
        "币种": "currency", "面值": "denomination", "张数": "quantity",
        "原币金额": "fcAmount", "汇率": "fxRate", "折算人民币": "rmbAmount",
    },
    "E1-9": {
        "存单编号": "certNo", "开户银行": "bank", "存款人/户名": "depositor",
        "账号": "account", "存款类型": "certType", "币种": "currency",
        "存入日": "depositDate", "到期日": "maturityDate", "金额": "amount",
        "利率(%)": "interestRate", "账面一致": "bookConsistent",
        "账面不一致原因": "inconsistencyReason", "是否质押/受限": "pledged",
        "质押/受限事项": "pledgeMatter", "盘点结果": "result",
        "存单/开户证实书索引": "certificateIndex", "开户证明索引": "openingProofIndex",
        "保管/质押证明索引": "custodyProofIndex", "备注": "note",
    },
    "E1-10": {
        "开户银行": "bank", "账号": "accountNo", "账户性质": "accountType",
        "开户日期": "openDate", "账户状态": "accountStatus", "销户日期": "closeDate",
        "开户原因": "openReason", "销户原因": "closeReason",
        "企业信息核对一致": "companyInfoConsistent", "企业信息不一致原因": "inconsistencyReason",
        "受限状态": "restrictionStatus", "开户目的(旧字段)": "openPurpose",
        "本期新开": "isNewThisPeriod", "本期注销": "isClosedThisPeriod",
        "账面有记录": "hasBookRecord", "清单核对一致": "checkResult", "差异说明": "reason",
    },
    "E1-15": {
        "存款类型": "depositType", "开户银行": "bank", "账号": "accountNo",
        "年利率": "annualRate",
        **{f"{month}月余额": f"balance{month}" for month in range(1, 13)},
        **{f"{month}月账面利息": f"bookInterest{month}" for month in range(1, 13)},
        **{f"{month}月测算利息": f"calculatedInterest{month}" for month in range(1, 13)},
    },
    "E1-20": {
        "开户银行": "bank", "账号": "accountNo", "用途": "usage", "类别": "category",
        "币种": "currency", "原币金额": "fcAmount", "结息日": "settleDate",
        "截止日": "cutoffDate", "天数": "days", "日利率": "dailyRate",
        "应计利息原币": "accruedFc", "汇率": "fxRate",
        "应计利息人民币": "accruedRmb", "备注": "note",
    },
    "E1-21": {
        "凭证号": "voucherNo", "日期": "date", "金额": "amount",
        "对方账户": "counterparty", "是否跨期": "isCrossover", "备注": "note",
    },
    "E1-22": {
        "凭证号": "voucherNo", "日期": "date", "金额": "amount",
        "对方账户": "counterparty", "是否跨期": "isCrossover", "备注": "note",
    },
}

_SUPPORTED_SHEETS = set(_FIELD_MAPS)
_SHEET_HEADERS = {sheet: list(fields) for sheet, fields in _FIELD_MAPS.items()}

_SHEET_ITEM_ID: dict[str, str] = {
    "E1-2": "E1-cash-detail-rows", "E1-3": "E1-bank-detail-rows",
    "E1-4": "E1-digital-rows", "E1-5": "E1-adjustment-rows",
    "E1-6": "E1-reconciliation-rows", "E1-7": "E1-cash-count-rmb-rows",
    "E1-8": "E1-cash-count-fx-rows", "E1-9": "E1-cash-count-cert-rows",
    "E1-10": "E1-account-list-rows", "E1-15": "E1-interest-monthly-rows",
    "E1-20": "E1-accrued-interest-rows", "E1-21": "E1-cutoff-bank-rows",
    "E1-22": "E1-cutoff-other-rows",
}


# 计算字段不属于前端持久化 JSON：导出时计算/回显，导入时忽略。
_FORMULA_FIELDS: dict[str, set[str]] = {
    "E1-2": {"endingFc", "endingRmb", "auditedRmb"},
    "E1-3": {"ending", "audited", "accountStatementDiff", "confirmDiff", "endingFc", "auditedFc"},
    "E1-4": {"endingFc", "endingRmb", "auditedFc", "auditedRmb", "diff"},
    "E1-5": set(),
    "E1-6": {"bankReceived", "bankPaid", "companyReceived", "companyPaid", "reconciledBook", "reconciledStatement", "diff"},
    "E1-7": {"subtotal"}, "E1-8": {"rmbAmount"}, "E1-9": set(), "E1-10": set(),
    "E1-15": {f"calculatedInterest{month}" for month in range(1, 13)},
    "E1-20": {"days", "accruedFc", "accruedRmb"},
    "E1-21": {"isCrossover"}, "E1-22": {"isCrossover"},
}

_NUMERIC_FIELDS: dict[str, set[str]] = {
    "E1-2": {"opening", "increase", "decrease", "endingFc", "fxRate", "endingRmb", "adjustment", "auditedRmb"},
    "E1-3": {"opening", "increase", "decrease", "ending", "adjustment", "audited", "statementBalance", "accountStatementDiff", "confirmAmount", "confirmDiff", "restrictedAmount", "interestRate", "fxRate", "openingFc", "increaseFc", "decreaseFc", "endingFc", "adjustmentFc", "auditedFc"},
    "E1-4": {"seq", "fxRate", "opening", "increase", "decrease", "endingFc", "endingRmb", "adjustment", "auditedFc", "auditedRmb", "queryBalance", "diff"},
    "E1-5": {"debit", "credit"},
    "E1-6": {"bookBalance", "statementBalance", "bankReceived", "bankPaid", "companyReceived", "companyPaid", "reconciledBook", "reconciledStatement", "diff"},
    "E1-7": {"denomination", "quantity", "subtotal"},
    "E1-8": {"denomination", "quantity", "fcAmount", "fxRate", "rmbAmount"},
    "E1-9": {"amount", "interestRate"}, "E1-10": set(),
    "E1-15": {"annualRate", *{f"balance{month}" for month in range(1, 13)}, *{f"bookInterest{month}" for month in range(1, 13)}, *{f"calculatedInterest{month}" for month in range(1, 13)}},
    "E1-20": {"fcAmount", "days", "dailyRate", "accruedFc", "fxRate", "accruedRmb"},
    "E1-21": {"amount"}, "E1-22": {"amount"},
}
_DATE_FIELDS = {
    "E1-9": {"depositDate", "maturityDate"}, "E1-10": {"openDate", "closeDate"},
    "E1-20": {"settleDate", "cutoffDate"}, "E1-21": {"date"}, "E1-22": {"date"},
}
_INTEGER_FIELDS = {"E1-4": {"seq"}}

_ENUM_MAPS: dict[tuple[str, str], dict[str, str]] = {
    ("E1-3", "section"): {
        "principal": "principal", "本金": "principal", "存款本金": "principal",
        "（一）存款本金": "principal", "accrued": "accrued", "应计利息": "accrued",
        "（二）应计利息": "accrued",
    },
    ("E1-3", "group"): {
        "principal": "institution", "存款本金": "institution",
        "institution": "institution", "银行机构": "institution",
        "其他金融机构": "institution", "金融机构": "institution",
        "finance": "finance", "财务公司": "finance", "存放财务公司款项": "finance",
        "其他金融机构（存放财务公司款项）": "finance",
        "other": "other", "其他货币资金": "other",
    },
    ("E1-5", "category"): {"报表调整": "报表调整", "账项调整": "账项调整", "其他": "其他"},
    ("E1-9", "bookConsistent"): {"是": "是", "Y": "是", "YES": "是", "否": "否", "N": "否", "NO": "否"},
    ("E1-9", "pledged"): {"是": "是", "Y": "是", "YES": "是", "否": "否", "N": "否", "NO": "否"},
    ("E1-9", "result"): {"已见": "已见", "未见": "未见"},
    ("E1-10", "companyInfoConsistent"): {"一致": "一致", "是": "一致", "不一致": "不一致", "否": "不一致"},
    ("E1-10", "isNewThisPeriod"): {"是": "Y", "Y": "Y", "YES": "Y", "否": "N", "N": "N", "NO": "N"},
    ("E1-10", "isClosedThisPeriod"): {"是": "Y", "Y": "Y", "YES": "Y", "否": "N", "N": "N", "NO": "N"},
    ("E1-10", "hasBookRecord"): {"是": "Y", "Y": "Y", "YES": "Y", "否": "N", "N": "N", "NO": "N"},
    ("E1-10", "checkResult"): {"一致": "一致", "是": "一致", "不一致": "不一致", "否": "不一致"},
    ("E1-20", "category"): {
        "finance": "finance", "财务公司": "finance", "bank": "bank", "银行": "bank",
        "other": "other", "其他": "other", "digital": "digital", "数字货币": "digital",
    },
}
_ENUM_DEFAULTS = {
    ("E1-3", "section"): "principal", ("E1-3", "group"): "institution",
    ("E1-5", "category"): "账项调整", ("E1-20", "category"): "bank",
}

# 旧后端字段只用于读取既有 remark 后导出；导入始终生成 canonical 字段。
_LEGACY_EXPORT_ALIASES: dict[str, dict[str, tuple[str, ...]]] = {
    "E1-2": {"opening": ("openingBalance",), "endingFc": ("endingBalance",), "fxRate": ("exchangeRate",), "endingRmb": ("convertedRmb",), "auditedRmb": ("audited",), "note": ("remark",)},
    "E1-3": {"opening": ("openingBalance",), "ending": ("endingBalance",), "accountStatementDiff": ("difference",), "confirmIndexNo": ("confirmIndexRef",)},
    "E1-5": {"debit": ("debitAmount",), "credit": ("creditAmount",), "indexNo": ("indexRef",), "note": ("remark",)},
    "E1-6": {"bankReceived": ("bookReceivedNotBank",), "bankPaid": ("bookPaidNotBank",), "companyReceived": ("bankReceivedNotBooked",), "companyPaid": ("bankPaidNotBooked",), "diff": ("difference",)},
    "E1-7": {"subtotal": ("amount",)}, "E1-8": {"fcAmount": ("amount",), "fxRate": ("exchangeRate",), "rmbAmount": ("convertedRmb",)},
    "E1-9": {"bank": ("bankName",), "note": ("remark",)},
    "E1-20": {"bank": ("bankName",), "fcAmount": ("principalAmount",), "settleDate": ("startDate",), "cutoffDate": ("maturityDate",), "fxRate": ("exchangeRate",), "accruedFc": ("accruedInterestFc",), "accruedRmb": ("accruedInterestRmb",), "note": ("remark",)},
    "E1-21": {"counterparty": ("summary",), "amount": ("receiptAmount",), "isCrossover": ("isCrossperiod",), "note": ("remark",)},
    "E1-22": {"counterparty": ("summary",), "amount": ("receiptAmount",), "isCrossover": ("isCrossperiod",), "note": ("remark",)},
}


def _validate_sheet(sheet_code: str) -> None:
    if sheet_code not in _SUPPORTED_SHEETS:
        raise HTTPException(400, f"不支持的sheet: {sheet_code}。支持: {sorted(_SUPPORTED_SHEETS)}")


def _safe_float(value: Any) -> float:
    if value is None or value == "":
        return 0.0
    try:
        return float(value)
    except (ValueError, TypeError):
        return 0.0


def _safe_str(value: Any) -> str:
    return "" if value is None else str(value).strip()


def _normalize_date(value: Any) -> str:
    if value is None or value == "":
        return ""
    if isinstance(value, datetime):
        return value.date().isoformat()
    if isinstance(value, date):
        return value.isoformat()
    if isinstance(value, (int, float)):
        try:
            return from_excel(value).date().isoformat()
        except (TypeError, ValueError, OverflowError):
            return _safe_str(value)
    text = _safe_str(value)
    normalized = text.replace("年", "-").replace("月", "-").replace("日", "").replace("/", "-").replace(".", "-")
    try:
        return datetime.fromisoformat(normalized).date().isoformat()
    except ValueError:
        return text


def _normalize_enum(sheet_code: str, field: str, value: Any) -> str:
    text = _safe_str(value)
    mapping = _ENUM_MAPS[(sheet_code, field)]
    return mapping.get(text, mapping.get(text.upper(), _ENUM_DEFAULTS.get((sheet_code, field), "")))


def _create_template_wb(sheet_code: str) -> Workbook:
    wb = Workbook()
    ws = wb.active
    ws.title = sheet_code
    font = Font(bold=True, size=11)
    fill = PatternFill(start_color="D9E1F2", end_color="D9E1F2", fill_type="solid")
    align = Alignment(horizontal="center", vertical="center", wrap_text=True)
    for col_idx, name in enumerate(_SHEET_HEADERS[sheet_code], 1):
        cell = ws.cell(row=1, column=col_idx, value=name)
        cell.font, cell.fill, cell.alignment = font, fill, align
        ws.column_dimensions[cell.column_letter].width = max(len(name) * 2 + 4, 12)
    ws.freeze_panes = "A2"
    return wb


def _get_actual_headers(ws: Any) -> list[str]:
    return [str(cell.value).strip() for cell in next(ws.iter_rows(min_row=1, max_row=1)) if cell.value is not None]


def _validate_columns(ws: Any, sheet_code: str) -> list[str]:
    actual = set(_get_actual_headers(ws))
    return [name for name in _SHEET_HEADERS[sheet_code] if name not in actual]


def _col_val(row: tuple, headers: list[str], name: str) -> Any:
    try:
        index = headers.index(name)
        return row[index] if index < len(row) else None
    except (ValueError, IndexError):
        return None


def _compat_value(sheet_code: str, data: dict[str, Any], field: str) -> Any:
    if field in data:
        return data[field]
    for alias in _LEGACY_EXPORT_ALIASES.get(sheet_code, {}).get(field, ()):
        if alias in data:
            return data[alias]
    return ""


def _days_between(start: Any, end: Any) -> int:
    start_text, end_text = _normalize_date(start), _normalize_date(end)
    try:
        return max(0, (date.fromisoformat(end_text) - date.fromisoformat(start_text)).days)
    except ValueError:
        return 0


def _sum_outstanding_items(value: Any) -> float:
    if not isinstance(value, list):
        return 0.0
    return sum(_safe_float(item.get("amount")) for item in value if isinstance(item, dict))


def _parse_outstanding_items(raw: Any, column: str) -> list[dict[str, Any]]:
    if raw is None or raw == "":
        return []
    try:
        loaded = json.loads(raw) if isinstance(raw, str) else raw
    except (json.JSONDecodeError, TypeError) as exc:
        raise HTTPException(400, f"{column}不是有效JSON") from exc
    if not isinstance(loaded, list) or any(not isinstance(item, dict) for item in loaded):
        raise HTTPException(400, f"{column}必须是JSON对象数组")
    return loaded


def _with_formula_values(sheet_code: str, data: dict[str, Any]) -> dict[str, Any]:
    if sheet_code == "E1-15":
        balances = data.get("balances") if isinstance(data.get("balances"), list) else []
        book_interests = data.get("bookInterests") if isinstance(data.get("bookInterests"), list) else []
        annual_rate = _safe_float(data.get("annualRate"))
        values: dict[str, Any] = {
            "depositType": _safe_str(data.get("depositType")),
            "bank": _safe_str(data.get("bank")),
            "accountNo": _safe_str(data.get("accountNo")),
            "annualRate": annual_rate,
        }
        for month in range(1, 13):
            balance = _safe_float(balances[month - 1] if month <= len(balances) else 0)
            values[f"balance{month}"] = balance
            values[f"bookInterest{month}"] = _safe_float(
                book_interests[month - 1] if month <= len(book_interests) else 0
            )
            values[f"calculatedInterest{month}"] = balance * annual_rate / 12
        return values

    values = {field: _compat_value(sheet_code, data, field) for field in _FIELD_MAPS[sheet_code].values()}
    n = lambda field: _safe_float(values.get(field))
    if sheet_code == "E1-2":
        values["endingFc"] = n("opening") + n("increase") - n("decrease")
        values["endingRmb"] = n("endingFc") * n("fxRate")
        values["auditedRmb"] = (n("endingFc") + n("adjustment")) * n("fxRate")
    elif sheet_code == "E1-3":
        values["endingFc"] = n("openingFc") + n("increaseFc") - n("decreaseFc")
        values["auditedFc"] = n("endingFc") + n("adjustmentFc")
        # opening/increase/decrease/adjustment 均是前端持久化字段，导出不得用
        # 外币列反向覆盖；这里只计算真正的公式字段。
        values["ending"] = n("opening") + n("increase") - n("decrease")
        values["audited"] = n("ending") + n("adjustment")
        values["accountStatementDiff"] = n("audited") - n("statementBalance")
        values["confirmDiff"] = n("audited") - n("confirmAmount") if n("confirmAmount") else 0.0
    elif sheet_code == "E1-4":
        values["endingFc"] = n("opening") + n("increase") - n("decrease")
        values["endingRmb"] = n("endingFc") * n("fxRate")
        values["auditedFc"] = n("endingFc") + n("adjustment")
        values["auditedRmb"] = n("auditedFc") * n("fxRate")
        values["diff"] = n("queryBalance") - n("auditedFc")
    elif sheet_code == "E1-6":
        category_totals = {
            "companyReceived": ("bankReceivedCompanyUnreceivedItems", "companyReceived"),
            "companyPaid": ("bankPaidCompanyUnpaidItems", "companyPaid"),
            "bankReceived": ("companyReceivedBankUnreceivedItems", "bankReceived"),
            "bankPaid": ("companyPaidBankUnpaidItems", "bankPaid"),
        }
        for total_field, (items_field, legacy_total_field) in category_totals.items():
            items = values.get(items_field)
            values[total_field] = (
                _sum_outstanding_items(items) if isinstance(items, list) else _safe_float(data.get(legacy_total_field))
            )
        values["reconciledBook"] = n("bookBalance") + n("companyReceived") - n("companyPaid")
        values["reconciledStatement"] = n("statementBalance") + n("bankReceived") - n("bankPaid")
        values["diff"] = n("reconciledBook") - n("reconciledStatement")
    elif sheet_code == "E1-7":
        values["subtotal"] = n("denomination") * n("quantity")
    elif sheet_code == "E1-8":
        values["rmbAmount"] = n("fcAmount") * n("fxRate")
    elif sheet_code == "E1-20":
        values["days"] = _days_between(values["settleDate"], values["cutoffDate"])
        values["accruedFc"] = n("fcAmount") * n("days") * n("dailyRate")
        values["accruedRmb"] = n("accruedFc") * n("fxRate")
    return values


def _export_row(sheet_code: str, data: dict[str, Any]) -> list[Any]:
    values = _with_formula_values(sheet_code, data)
    exported: list[Any] = []
    for field in _FIELD_MAPS[sheet_code].values():
        value = values.get(field, "")
        if sheet_code == "E1-6" and field in {
            "bankReceivedCompanyUnreceivedItems", "bankPaidCompanyUnpaidItems",
            "companyReceivedBankUnreceivedItems", "companyPaidBankUnpaidItems",
        }:
            value = json.dumps(value if isinstance(value, list) else [], ensure_ascii=False, separators=(",", ":"))
        exported.append(value if value is not None else "")
    return exported


def _parse_row(sheet_code: str, row: tuple, actual_headers: list[str]) -> dict[str, Any]:
    result: dict[str, Any] = {"id": str(uuid4())}
    if sheet_code == "E1-15":
        result.update({
            "depositType": _safe_str(_col_val(row, actual_headers, "存款类型")),
            "bank": _safe_str(_col_val(row, actual_headers, "开户银行")),
            "accountNo": _safe_str(_col_val(row, actual_headers, "账号")),
            "annualRate": _safe_float(_col_val(row, actual_headers, "年利率")),
            "balances": [
                _safe_float(_col_val(row, actual_headers, f"{month}月余额"))
                for month in range(1, 13)
            ],
            "bookInterests": [
                _safe_float(_col_val(row, actual_headers, f"{month}月账面利息"))
                for month in range(1, 13)
            ],
        })
        return result

    formulas = _FORMULA_FIELDS[sheet_code]
    nested_fields = {
        "bankReceivedCompanyUnreceivedItems", "bankPaidCompanyUnpaidItems",
        "companyReceivedBankUnreceivedItems", "companyPaidBankUnpaidItems",
    }
    for column, field in _FIELD_MAPS[sheet_code].items():
        if field in formulas:
            continue
        raw = _col_val(row, actual_headers, column)
        if sheet_code == "E1-6" and field in nested_fields:
            value: Any = _parse_outstanding_items(raw, column)
        elif (sheet_code, field) in _ENUM_MAPS:
            value = _normalize_enum(sheet_code, field, raw)
        elif field in _DATE_FIELDS.get(sheet_code, set()):
            value = _normalize_date(raw)
        elif field in _NUMERIC_FIELDS[sheet_code]:
            value = _safe_float(raw)
            if field in _INTEGER_FIELDS.get(sheet_code, set()):
                value = int(value)
        else:
            value = _safe_str(raw)
        result[field] = value
    return result


def _xlsx_response(wb: Workbook, filename: str) -> StreamingResponse:
    buffer = io.BytesIO()
    wb.save(buffer)
    buffer.seek(0)
    return StreamingResponse(
        buffer,
        media_type="application/vnd.openxmlformats-officedocument.spreadsheetml.sheet",
        headers={"Content-Disposition": f"attachment; filename*=UTF-8''{quote(filename)}"},
    )


@router.post("/api/workpapers/{wp_id}/e1/export-template")
async def e1_export_template(
    wp_id: str,
    sheet: str = Query(..., description="E1 sheet 编码"),
    current_user: User = Depends(get_current_user),
) -> StreamingResponse:
    _validate_sheet(sheet)
    return _xlsx_response(_create_template_wb(sheet), f"{sheet}_模板.xlsx")


@router.post("/api/workpapers/{wp_id}/e1/export-data")
async def e1_export_data(
    wp_id: str,
    sheet: str = Query(..., description="E1 sheet 编码"),
    db: AsyncSession = Depends(get_db),
    current_user: User = Depends(get_current_user),
) -> StreamingResponse:
    _validate_sheet(sheet)
    item_id = _SHEET_ITEM_ID[sheet]
    result = await db.execute(
        sa.text("SELECT remark FROM checklist_responses WHERE wp_id = :wp_id AND item_id = :item_id LIMIT 1"),
        {"wp_id": wp_id, "item_id": item_id},
    )
    stored = result.fetchone()
    rows_data: list[dict[str, Any]] = []
    if stored and stored.remark:
        try:
            loaded = json.loads(stored.remark)
            if isinstance(loaded, list):
                rows_data = [row for row in loaded if isinstance(row, dict)]
        except (json.JSONDecodeError, TypeError):
            logger.warning("E1 export ignored invalid JSON: wp=%s item=%s", wp_id, item_id)

    wb = _create_template_wb(sheet)
    for data_row in rows_data:
        wb.active.append(_export_row(sheet, data_row))
    return _xlsx_response(wb, f"{sheet}_数据.xlsx")


def _project_id_from_row(row: Any) -> Any:
    if row is None:
        return None
    if hasattr(row, "project_id"):
        return row.project_id
    try:
        return row[0]
    except (TypeError, IndexError, KeyError):
        return None


@router.post("/api/workpapers/{wp_id}/e1/import-data")
async def e1_import_data(
    wp_id: str,
    sheet: str = Query(..., description="E1 sheet 编码"),
    file: UploadFile = File(...),
    db: AsyncSession = Depends(get_db),
    current_user: User = Depends(get_current_user),
) -> dict[str, Any]:
    _validate_sheet(sheet)
    content = await file.read()
    try:
        wb = load_workbook(io.BytesIO(content), read_only=True, data_only=True)
    except Exception as exc:
        raise HTTPException(400, "无法解析xlsx文件") from exc

    ws = wb.active
    missing = _validate_columns(ws, sheet)
    if missing:
        raise HTTPException(400, f"列名不匹配，缺少: {missing}")
    headers = _get_actual_headers(ws)
    parsed_rows: list[dict[str, Any]] = []
    for row_idx, row in enumerate(ws.iter_rows(min_row=2, values_only=True), start=2):
        if row_idx > _ROW_LIMIT + 1:
            break
        if all(cell is None or str(cell).strip() == "" for cell in row):
            continue
        parsed_rows.append(_parse_row(sheet, row, headers))

    # checklist_responses.project_id NOT NULL；即使 ON CONFLICT 命中也会先校验 INSERT。
    project_result = await db.execute(
        sa.text("SELECT project_id FROM working_paper WHERE id = :wp_id"),
        {"wp_id": wp_id},
    )
    project_id = _project_id_from_row(project_result.fetchone())
    if project_id is None:
        raise HTTPException(404, "底稿不存在或未关联项目")

    item_id = _SHEET_ITEM_ID[sheet]
    params = {
        "id": str(uuid4()), "project_id": project_id, "wp_id": wp_id,
        "item_id": item_id, "remark": json.dumps(parsed_rows, ensure_ascii=False),
    }
    await db.execute(
        sa.text(
            "INSERT INTO checklist_responses (id, project_id, wp_id, item_id, remark) "
            "VALUES (:id, :project_id, :wp_id, :item_id, :remark) "
            "ON CONFLICT (wp_id, item_id) DO UPDATE "
            "SET project_id = EXCLUDED.project_id, remark = EXCLUDED.remark"
        ),
        params,
    )
    await db.commit()
    return {"row_count": len(parsed_rows), "sheet": sheet, "item_id": item_id}

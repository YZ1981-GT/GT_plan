"""H1-9 固定资产监盘计划 — Word 导出

POST /api/workpapers/{wp_id}/h1/stocktake-plan-export
Body: { fields: H1StocktakePlanForm-like dict, entity_name?: str }
返回 docx 文件流（Content-Disposition: attachment）
"""

from __future__ import annotations

import io
import logging
from typing import Any

from fastapi import APIRouter, Depends, HTTPException
from fastapi.responses import StreamingResponse
from pydantic import BaseModel, Field
from sqlalchemy.ext.asyncio import AsyncSession

from app.core.database import get_db
from app.deps import get_current_user
from app.models.core import User

logger = logging.getLogger(__name__)

router = APIRouter(tags=["h1-ai"])


class H1StocktakePlanExportRequest(BaseModel):
    fields: dict[str, Any] = Field(default_factory=dict)
    entity_name: str = ""
    index_no: str = "H1-9"


def _txt(v: Any) -> str:
    if v is None:
        return ""
    if isinstance(v, (int, float)):
        return str(v)
    return str(v).strip()


def build_plan_docx(fields: dict[str, Any], entity_name: str = "", index_no: str = "H1-9") -> bytes:
    try:
        from docx import Document
        from docx.enum.text import WD_ALIGN_PARAGRAPH
        from docx.shared import Pt
    except ImportError as e:
        raise RuntimeError("python-docx 未安装") from e

    doc = Document()
    title = doc.add_paragraph()
    run = title.add_run("固定资产监盘计划")
    run.bold = True
    run.font.size = Pt(16)
    title.alignment = WD_ALIGN_PARAGRAPH.CENTER

    meta = doc.add_paragraph()
    meta.add_run(f"被审计单位：{entity_name or '________'}    索引号：{index_no}")

    def h(text: str) -> None:
        p = doc.add_paragraph()
        r = p.add_run(text)
        r.bold = True
        r.font.size = Pt(12)

    def line(label: str, value: Any) -> None:
        doc.add_paragraph(f"{label}{_txt(value) or '________'}")

    h("一、固定资产存在性认定重大错报风险")
    line("风险程度：", fields.get("existenceRiskLevel"))
    line("评估说明：", fields.get("existenceRiskNote"))

    h("二、了解固定资产期末状况与管理规定")
    line("（一）闲置资产：", fields.get("idleAssetsNote"))
    locs = fields.get("locations") or []
    if isinstance(locs, list) and locs:
        doc.add_paragraph("存放地点：")
        for i, loc in enumerate(locs, 1):
            if not isinstance(loc, dict):
                continue
            doc.add_paragraph(
                f"  {i}. {loc.get('category', '')} / {loc.get('warehouse', '')} / "
                f"{loc.get('place', '')} — {loc.get('note', '')}"
            )
    line("制度名称：", fields.get("icSystemName"))
    line("制度索引：", fields.get("icSystemIndex"))
    line("盘点频次：", fields.get("icFrequency"))
    line("负责部门/人员：", fields.get("icResponsible"))
    line("企业盘点安排：", fields.get("clientPlanArrange"))
    line("会议信息：", fields.get("clientMeetingNote"))
    line("预计人数：", fields.get("clientHeadcountEstimate"))
    line("对盘点计划评价：", fields.get("clientPlanEvaluation"))
    line("上年监盘日期：", fields.get("priorYearDate"))
    line("上年人员：", fields.get("priorYearStaff"))
    line("上年范围：", fields.get("priorYearScope"))
    line("上年抽样比例：", fields.get("priorYearSampleRate"))
    line("上年问题与不足：", fields.get("priorYearIssues"))

    h("三、评估审计人员的专业胜任能力")
    line("", fields.get("competenceNote"))

    h("四、盘点计划安排")
    line("预计监盘日期：", fields.get("plannedDate"))
    line("时段说明：", fields.get("plannedTimeNote"))
    line("预计人数：", fields.get("plannedHeadcount"))
    line("程序负责人：", fields.get("plannedLead"))
    scopes = fields.get("categoryScopes") or []
    if isinstance(scopes, list) and scopes:
        doc.add_paragraph("类别范围：")
        table = doc.add_table(rows=1, cols=6)
        hdr = table.rows[0].cells
        for i, name in enumerate(["类别", "期末余额", "净值", "计划数量", "计划金额", "比例%"]):
            hdr[i].text = name
        for row in scopes:
            if not isinstance(row, dict):
                continue
            cells = table.add_row().cells
            cells[0].text = _txt(row.get("category"))
            cells[1].text = _txt(row.get("endingBalance"))
            cells[2].text = _txt(row.get("netBookValue"))
            cells[3].text = _txt(row.get("planQty"))
            cells[4].text = _txt(row.get("planAmount"))
            cells[5].text = _txt(row.get("coverageRate"))
    line("地点范围：", fields.get("locationScopeNote"))
    line("盘点方法：", fields.get("method"))
    line("方法说明：", fields.get("methodDetail"))
    line("与管理层沟通时间：", fields.get("mgmtCommTime"))
    line("管理层人员：", fields.get("mgmtCommNames"))
    line("审计人员：", fields.get("mgmtCommAuditors"))
    line("特殊要求：", fields.get("specialRequirements"))
    line("账面→实物：", f"{fields.get('sampleBookToFloorMethod', '')} 预计 {fields.get('sampleBookToFloorQty', '')}")
    line("实物→账面：", f"{fields.get('sampleFloorToBookMethod', '')} 预计 {fields.get('sampleFloorToBookQty', '')}")
    line("推算方法：", fields.get("rollForwardMethod"))
    line("预计复盘比例(%)：", fields.get("plannedRecountRatio"))

    h("五、监盘计划结论")
    line("", fields.get("planConclusion"))
    line("编制人：", f"{fields.get('preparedBy', '')}  日期：{fields.get('preparedDate', '')}")
    line("复核人：", f"{fields.get('reviewedBy', '')}  日期：{fields.get('reviewedDate', '')}")

    buf = io.BytesIO()
    doc.save(buf)
    return buf.getvalue()


@router.post("/api/workpapers/{wp_id}/h1/stocktake-plan-export")
async def h1_stocktake_plan_export(
    wp_id: str,
    body: H1StocktakePlanExportRequest,
    db: AsyncSession = Depends(get_db),
    current_user: User = Depends(get_current_user),
):
    _ = db, current_user, wp_id
    try:
        data = build_plan_docx(body.fields, body.entity_name, body.index_no)
    except RuntimeError as e:
        raise HTTPException(503, str(e)) from e
    except Exception as e:
        logger.exception("H1-9 Word export failed")
        raise HTTPException(500, f"导出失败: {e}") from e

    filename = "H1-9_固定资产监盘计划.docx"
    return StreamingResponse(
        io.BytesIO(data),
        media_type="application/vnd.openxmlformats-officedocument.wordprocessingml.document",
        headers={"Content-Disposition": f'attachment; filename="{filename}"'},
    )

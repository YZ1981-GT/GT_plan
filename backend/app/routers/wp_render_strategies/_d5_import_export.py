"""D5 应收款项融资 — 导入导出端点

4个端点：
- POST /api/workpapers/{wp_id}/d5/export-template?sheet={sheet_code}  空白模板xlsx
- POST /api/workpapers/{wp_id}/d5/export-data?sheet={sheet_code}      数据xlsx
- POST /api/workpapers/{wp_id}/d5/import-data?sheet={sheet_code}      解析xlsx写入
- POST /api/workpapers/{wp_id}/d5/import-aux-balance                  从辅助余额表导入

支持sheets: D5-2, D5-4
"""

from __future__ import annotations

import io
import json
import logging
from typing import Any
from uuid import uuid4

from fastapi import APIRouter, Depends, File, HTTPException, Query, UploadFile
from fastapi.responses import StreamingResponse
from openpyxl import Workbook, load_workbook
from sqlalchemy.ext.asyncio import AsyncSession

from app.core.database import get_db
from app.deps import get_current_user
from app.models.core import User

logger = logging.getLogger(__name__)

router = APIRouter(tags=["d5-import-export"])

# ═══════════════════════════════════════════════════════════════════════════════
# Sheet 配置
# ═══════════════════════════════════════════════════════════════════════════════

_ROW_LIMIT = 500

_SUPPORTED_SHEETS: set[str] = {"D5-2", "D5-4"}

_SHEET_HEADERS: dict[str, list[str]] = {
    "D5-2": [
        "类别", "明细项目", "期初未审", "期初AJE", "期初RJE", "期初审定",
        "OCI减值", "本期增加", "本期减少", "期末余额", "重分类",
        "期末未审", "期末AJE", "期末RJE", "期末审定", "期末OCI减值", "备注",
    ],
    "D5-4": [
        "类别", "明细项目", "票据号", "票面金额", "计量日", "到期日",
        "剩余天数", "贴现利率", "贴现利息", "贴现金额", "公允价值", "层次", "备注",
    ],
}

# item_id 映射
_SHEET_ITEM_ID: dict[str, str] = {
    "D5-2": "D5-2-rows",
    "D5-4": "D5-4-rows",
}


def _get_headers(sheet_code: str) -> list[str]:
    return _SHEET_HEADERS.get(sheet_code, [])


def _validate_sheet(sheet_code: str) -> None:
    if sheet_code not in _SUPPORTED_SHEETS:
        raise HTTPException(
            400,
            f"不支持的sheet: {sheet_code}。支持: {sorted(_SUPPORTED_SHEETS)}",
        )


# ═══════════════════════════════════════════════════════════════════════════════
# 端点
# ═══════════════════════════════════════════════════════════════════════════════


@router.post("/api/workpapers/{wp_id}/d5/export-template")
async def d5_export_template(
    wp_id: str,
    sheet: str = Query(..., description="Sheet编码如D5-2"),
    include_guidance: bool = Query(True, description="是否包含编制说明"),
    current_user: User = Depends(get_current_user),
) -> StreamingResponse:
    """导出空白模板xlsx（含表头+格式+编制说明，无数据行）"""
    _validate_sheet(sheet)
    headers = _get_headers(sheet)

    wb = Workbook()
    ws = wb.active
    ws.title = sheet
    ws.append(headers)
    ws.freeze_panes = "A2"
    for col_idx in range(1, len(headers) + 1):
        ws.column_dimensions[chr(64 + min(col_idx, 26))].width = 16

    # 添加编制说明 sheet
    if include_guidance:
        guidance = _get_d5_guidance_text(sheet)
        if guidance:
            ws_guide = wb.create_sheet("编制说明")
            ws_guide.append(["编制说明"])
            ws_guide.append([])
            for line in guidance:
                ws_guide.append([line])
            ws_guide.column_dimensions["A"].width = 80

    buffer = io.BytesIO()
    wb.save(buffer)
    buffer.seek(0)

    filename = f"{sheet}_模板.xlsx"
    from urllib.parse import quote
    encoded_filename = quote(filename)
    return StreamingResponse(
        buffer,
        media_type="application/vnd.openxmlformats-officedocument.spreadsheetml.sheet",
        headers={"Content-Disposition": f"attachment; filename*=UTF-8''{encoded_filename}"},
    )


@router.post("/api/workpapers/{wp_id}/d5/export-data")
async def d5_export_data(
    wp_id: str,
    sheet: str = Query(..., description="Sheet编码如D5-2"),
    db: AsyncSession = Depends(get_db),
    current_user: User = Depends(get_current_user),
) -> StreamingResponse:
    """导出当前数据xlsx"""
    _validate_sheet(sheet)
    headers = _get_headers(sheet)

    import sqlalchemy as sa

    item_id = _SHEET_ITEM_ID[sheet]
    result = await db.execute(
        sa.text(
            "SELECT remark FROM checklist_responses "
            "WHERE wp_id = :wp_id AND item_id = :item_id LIMIT 1"
        ),
        {"wp_id": wp_id, "item_id": item_id},
    )
    row = result.fetchone()
    rows_data: list[dict] = []
    if row and row.remark:
        try:
            rows_data = json.loads(row.remark)
        except (json.JSONDecodeError, TypeError):
            pass

    wb = Workbook()
    ws = wb.active
    ws.title = sheet
    ws.append(headers)
    ws.freeze_panes = "A2"

    for data_row in rows_data:
        if sheet == "D5-2":
            row_values = _export_d5_2_row(data_row)
        else:  # D5-4
            row_values = _export_d5_4_row(data_row)
        ws.append(row_values)

    buffer = io.BytesIO()
    wb.save(buffer)
    buffer.seek(0)

    filename = f"{sheet}_数据.xlsx"
    from urllib.parse import quote
    encoded_filename = quote(filename)
    return StreamingResponse(
        buffer,
        media_type="application/vnd.openxmlformats-officedocument.spreadsheetml.sheet",
        headers={"Content-Disposition": f"attachment; filename*=UTF-8''{encoded_filename}"},
    )


@router.post("/api/workpapers/{wp_id}/d5/import-data")
async def d5_import_data(
    wp_id: str,
    sheet: str = Query(..., description="Sheet编码如D5-2"),
    file: UploadFile = File(...),
    db: AsyncSession = Depends(get_db),
    current_user: User = Depends(get_current_user),
) -> dict[str, Any]:
    """解析上传xlsx，校验格式，写入 checklist_responses"""
    _validate_sheet(sheet)

    if not file.filename or not file.filename.endswith(".xlsx"):
        raise HTTPException(400, "请上传 .xlsx 格式文件")

    content = await file.read()
    if len(content) > 10 * 1024 * 1024:
        raise HTTPException(400, "文件大小不能超过10MB")

    try:
        wb = load_workbook(io.BytesIO(content), read_only=True, data_only=True)
    except Exception:
        raise HTTPException(400, "无法解析xlsx文件，请确认文件格式正确")

    ws = wb.active
    if ws is None:
        raise HTTPException(400, "xlsx文件中无活动工作表")

    # 验证列头
    expected_headers = _get_headers(sheet)
    actual_headers = [
        str(cell.value).strip() if cell.value else ""
        for cell in next(ws.iter_rows(min_row=1, max_row=1))
    ]

    errors: list[str] = []
    missing_cols = [h for h in expected_headers if h not in actual_headers]
    if missing_cols:
        errors.append(f"缺少列: {', '.join(missing_cols)}")

    if errors:
        raise HTTPException(400, detail=errors)

    # 解析数据行
    rows_data: list[dict] = []
    truncated = False
    row_count = 0

    for row in ws.iter_rows(min_row=2, values_only=True):
        if all(v is None for v in row):
            continue
        row_count += 1
        if row_count > _ROW_LIMIT:
            truncated = True
            break

        if sheet == "D5-2":
            row_dict = _parse_d5_2_row(row, actual_headers)
        else:
            row_dict = _parse_d5_4_row(row, actual_headers)

        rows_data.append(row_dict)

    wb.close()

    # 写入 checklist_responses
    import sqlalchemy as sa

    item_id = _SHEET_ITEM_ID[sheet]
    remark_json = json.dumps(rows_data, ensure_ascii=False)

    await db.execute(
        sa.text("""
            INSERT INTO checklist_responses (id, wp_id, item_id, remark, updated_at)
            VALUES (:id, :wp_id, :item_id, :remark, NOW())
            ON CONFLICT (wp_id, item_id)
            DO UPDATE SET remark = :remark, updated_at = NOW()
        """),
        {"id": str(uuid4()), "wp_id": wp_id, "item_id": item_id, "remark": remark_json},
    )
    await db.commit()

    result_data: dict[str, Any] = {"ok": True, "imported_count": len(rows_data), "errors": []}
    if truncated:
        result_data["warning"] = f"数据行数超过{_ROW_LIMIT}行限制，已截断"
        result_data["truncated"] = True

    return result_data


@router.post("/api/workpapers/{wp_id}/d5/import-aux-balance")
async def d5_import_aux_balance(
    wp_id: str,
    db: AsyncSession = Depends(get_db),
    current_user: User = Depends(get_current_user),
) -> dict[str, Any]:
    """从 tb_aux_balance（科目1124，按客户维度）批量导入到D5-2"""
    import sqlalchemy as sa

    # 获取 project_id
    wp_result = await db.execute(
        sa.text("SELECT project_id FROM working_paper WHERE id = :wp_id"),
        {"wp_id": wp_id},
    )
    wp_row = wp_result.fetchone()
    if not wp_row:
        raise HTTPException(404, "底稿不存在")

    project_id = str(wp_row.project_id)

    # 从 tb_aux_balance 查询科目1124按客户维度聚合
    aux_result = await db.execute(
        sa.text("""
            SELECT aux_name,
                   COALESCE(SUM(CASE WHEN period_type = 'opening' THEN balance ELSE 0 END), 0) AS prior_balance,
                   COALESCE(SUM(CASE WHEN period_type = 'closing' THEN balance ELSE 0 END), 0) AS current_balance
            FROM tb_aux_balance
            WHERE project_id = :pid
              AND account_code = '1124'
              AND is_deleted = false
            GROUP BY aux_name
            ORDER BY aux_name
        """),
        {"pid": project_id},
    )
    aux_rows = aux_result.fetchall()

    if not aux_rows:
        return {"ok": True, "imported_count": 0, "message": "未找到科目1124的辅助余额数据"}

    # 构建 D5-2 行数据
    rows_data: list[dict] = []
    for aux_row in aux_rows[:_ROW_LIMIT]:
        rows_data.append({
            "rowId": str(uuid4()),
            "category": "应收票据",
            "itemName": aux_row.aux_name or "",
            "priorUnadjusted": float(aux_row.prior_balance),
            "priorAje": 0,
            "priorRje": 0,
            "priorAudited": float(aux_row.prior_balance),
            "ociImpairment": 0,
            "periodIncrease": 0,
            "periodDecrease": 0,
            "endBalance": float(aux_row.current_balance),
            "entityReclass": 0,
            "endUnadjusted": float(aux_row.current_balance),
            "endAje": 0,
            "endRje": 0,
            "endAudited": float(aux_row.current_balance),
            "endOciImpairment": 0,
            "remark": "",
        })

    # 写入（merge模式：保留已有行、追加新客户）
    item_id = "D5-2-rows"
    existing_result = await db.execute(
        sa.text("SELECT remark FROM checklist_responses WHERE wp_id = :wp_id AND item_id = :item_id LIMIT 1"),
        {"wp_id": wp_id, "item_id": item_id},
    )
    existing_row = existing_result.fetchone()
    existing_rows: list[dict] = []
    if existing_row and existing_row.remark:
        try:
            existing_rows = json.loads(existing_row.remark)
        except (json.JSONDecodeError, TypeError):
            pass

    # Merge: 已有明细项目不重复导入
    existing_names = {r.get("itemName", "") for r in existing_rows}
    new_rows = [r for r in rows_data if r["itemName"] not in existing_names]
    merged = existing_rows + new_rows

    remark_json = json.dumps(merged, ensure_ascii=False)
    await db.execute(
        sa.text("""
            INSERT INTO checklist_responses (id, wp_id, item_id, remark, updated_at)
            VALUES (:id, :wp_id, :item_id, :remark, NOW())
            ON CONFLICT (wp_id, item_id)
            DO UPDATE SET remark = :remark, updated_at = NOW()
        """),
        {"id": str(uuid4()), "wp_id": wp_id, "item_id": item_id, "remark": remark_json},
    )
    await db.commit()

    return {
        "ok": True,
        "imported_count": len(new_rows),
        "message": f"成功导入{len(new_rows)}行数据",
        "total_rows": len(merged),
    }


# ═══════════════════════════════════════════════════════════════════════════════
# 导出辅助函数
# ═══════════════════════════════════════════════════════════════════════════════


def _safe_float(val: Any) -> float:
    if val is None:
        return 0.0
    try:
        return float(val)
    except (ValueError, TypeError):
        return 0.0


def _safe_str(val: Any) -> str:
    if val is None:
        return ""
    return str(val).strip()


def _export_d5_2_row(data: dict) -> list:
    return [
        _safe_str(data.get("category")),
        _safe_str(data.get("itemName")),
        _safe_float(data.get("priorUnadjusted")),
        _safe_float(data.get("priorAje")),
        _safe_float(data.get("priorRje")),
        _safe_float(data.get("priorAudited")),
        _safe_float(data.get("ociImpairment")),
        _safe_float(data.get("periodIncrease")),
        _safe_float(data.get("periodDecrease")),
        _safe_float(data.get("endBalance")),
        _safe_float(data.get("entityReclass")),
        _safe_float(data.get("endUnadjusted")),
        _safe_float(data.get("endAje")),
        _safe_float(data.get("endRje")),
        _safe_float(data.get("endAudited")),
        _safe_float(data.get("endOciImpairment")),
        _safe_str(data.get("remark")),
    ]


def _export_d5_4_row(data: dict) -> list:
    return [
        _safe_str(data.get("category")),
        _safe_str(data.get("itemName")),
        _safe_str(data.get("billNo")),
        _safe_float(data.get("faceValue")),
        _safe_str(data.get("measurementDate")),
        _safe_str(data.get("maturityDate")),
        _safe_float(data.get("remainingDays")),
        _safe_float(data.get("discountRate")),
        _safe_float(data.get("discountInterest")),
        _safe_float(data.get("discountAmount")),
        _safe_float(data.get("fairValue")),
        _safe_str(data.get("fvHierarchy")),
        _safe_str(data.get("remark")),
    ]


# ═══════════════════════════════════════════════════════════════════════════════
# 导入解析辅助函数
# ═══════════════════════════════════════════════════════════════════════════════


def _col_val(row: tuple, actual_headers: list[str], col_name: str) -> Any:
    try:
        idx = actual_headers.index(col_name)
        return row[idx] if idx < len(row) else None
    except (ValueError, IndexError):
        return None


def _parse_d5_2_row(row: tuple, actual_headers: list[str]) -> dict:
    prior_unadj = _safe_float(_col_val(row, actual_headers, "期初未审"))
    prior_aje = _safe_float(_col_val(row, actual_headers, "期初AJE"))
    prior_rje = _safe_float(_col_val(row, actual_headers, "期初RJE"))
    prior_audited = _safe_float(_col_val(row, actual_headers, "期初审定"))
    # If prior_audited not explicitly set, compute from formula
    if prior_audited == 0.0 and (prior_unadj != 0.0 or prior_aje != 0.0 or prior_rje != 0.0):
        prior_audited = prior_unadj + prior_aje + prior_rje

    end_unadj = _safe_float(_col_val(row, actual_headers, "期末未审"))
    end_aje = _safe_float(_col_val(row, actual_headers, "期末AJE"))
    end_rje = _safe_float(_col_val(row, actual_headers, "期末RJE"))
    end_audited = _safe_float(_col_val(row, actual_headers, "期末审定"))
    if end_audited == 0.0 and (end_unadj != 0.0 or end_aje != 0.0 or end_rje != 0.0):
        end_audited = end_unadj + end_aje + end_rje

    return {
        "rowId": str(uuid4()),
        "category": _safe_str(_col_val(row, actual_headers, "类别")) or "应收票据",
        "itemName": _safe_str(_col_val(row, actual_headers, "明细项目")),
        "priorUnadjusted": prior_unadj,
        "priorAje": prior_aje,
        "priorRje": prior_rje,
        "priorAudited": prior_audited,
        "ociImpairment": _safe_float(_col_val(row, actual_headers, "OCI减值")),
        "periodIncrease": _safe_float(_col_val(row, actual_headers, "本期增加")),
        "periodDecrease": _safe_float(_col_val(row, actual_headers, "本期减少")),
        "endBalance": _safe_float(_col_val(row, actual_headers, "期末余额")),
        "entityReclass": _safe_float(_col_val(row, actual_headers, "重分类")),
        "endUnadjusted": end_unadj,
        "endAje": end_aje,
        "endRje": end_rje,
        "endAudited": end_audited,
        "endOciImpairment": _safe_float(_col_val(row, actual_headers, "期末OCI减值")),
        "remark": _safe_str(_col_val(row, actual_headers, "备注")),
    }


def _parse_d5_4_row(row: tuple, actual_headers: list[str]) -> dict:
    face_value = _safe_float(_col_val(row, actual_headers, "票面金额"))
    remaining_days = _safe_float(_col_val(row, actual_headers, "剩余天数"))
    discount_rate = _safe_float(_col_val(row, actual_headers, "贴现利率"))
    discount_interest = _safe_float(_col_val(row, actual_headers, "贴现利息"))
    # Compute discount_interest if not explicitly provided
    if discount_interest == 0.0 and face_value > 0 and discount_rate > 0 and remaining_days > 0:
        discount_interest = face_value * discount_rate * remaining_days / 360
    fair_value = _safe_float(_col_val(row, actual_headers, "公允价值"))
    if fair_value == 0.0 and face_value > 0:
        fair_value = face_value - discount_interest

    return {
        "rowId": str(uuid4()),
        "category": _safe_str(_col_val(row, actual_headers, "类别")) or "应收票据",
        "itemName": _safe_str(_col_val(row, actual_headers, "明细项目")),
        "billNo": _safe_str(_col_val(row, actual_headers, "票据号")),
        "faceValue": face_value,
        "measurementDate": _safe_str(_col_val(row, actual_headers, "计量日")),
        "maturityDate": _safe_str(_col_val(row, actual_headers, "到期日")),
        "remainingDays": remaining_days,
        "discountRate": discount_rate,
        "discountInterest": discount_interest,
        "discountAmount": _safe_float(_col_val(row, actual_headers, "贴现金额")) or (face_value - discount_interest),
        "fairValue": fair_value,
        "fvHierarchy": _safe_str(_col_val(row, actual_headers, "层次")) or "第二层次",
        "remark": _safe_str(_col_val(row, actual_headers, "备注")),
    }


# ═══════════════════════════════════════════════════════════════════════════════
# 编制说明文本（导出模板时附加）
# ═══════════════════════════════════════════════════════════════════════════════

_D5_SHEET_GUIDANCE: dict[str, list[str]] = {
    "D5-2": [
        "D5-2 应收款项融资明细表 编制说明",
        "",
        "一、本表目的",
        "按类别（应收票据/应收账款）列示以公允价值计量且其变动计入其他综合收益（FVOCI）的应收款项融资明细。",
        "反映各项资产的期初、本期变动和期末审定余额。",
        "",
        "二、科目特征",
        "科目编码：1124 应收款项融资（借方科目/资产类/FVOCI金融资产）",
        "业务模式：收取合同现金流量+出售（既收本息又可出售变现）",
        "",
        "三、填写要求",
        '1. 「类别」列：选择"应收票据"或"应收账款"（从D1/D2业务模式判定为出售模式的资产归入此科目）。',
        "2. 「明细项目」列：填写具体的票据出票人/应收账款客户名称。",
        "3. 「期初未审(C)」「期初AJE(D)」「期初RJE(E)」列：填写期初数据。",
        "4. 「OCI减值(G)」列：填写期初OCI减值准备余额。",
        "5. 「本期增加(H)」「本期减少(I)」列：填写本期发生额（借方增加/贷方减少）。",
        "6. 「重分类(K)」列：如被审计单位有重分类调整则填写。",
        "7. 「期末AJE(M)」「期末RJE(N)」列：审计调整分录金额。",
        "",
        "四、自动计算列（灰色背景，无需填写）",
        "F = C + D + E（期初审定 = 期初未审 + AJE + RJE）",
        "J = F + H - I（期末余额 = 期初审定 + 本期增加 - 本期减少）",
        "L = J + K（期末未审 = 期末余额 + 重分类）",
        "O = L + M + N（期末审定 = 期末未审 + 期末AJE + 期末RJE）",
        "",
        "五、审计关注",
        "1. 类别判定需与D1(票据)/D2(应收账款)的业务模式判定一致。",
        "2. 期末审定合计应与D5-1审定表中对应类别行一致。",
        "3. 本表数据将传递给D5-4公允价值测算（计算OCI公允价值变动）。",
    ],
    "D5-4": [
        "D5-4 应收款项融资公允价值测算表 编制说明",
        "",
        "一、本表目的",
        "对FVOCI金融资产（应收款项融资）进行期末公允价值测算，核心采用贴现法。",
        "",
        "二、核心公式",
        "剩余天数 = 到期日 - 计量日（天数）",
        "贴现利息 = 票面金额 × 市场贴现利率 × 剩余天数 ÷ 360",
        "公允价值 = 票面金额 - 贴现利息",
        "",
        "三、填写要求",
        '1. 「类别」列：选择"应收票据"或"应收账款"。',
        "2. 「明细项目」列：与D5-2明细表对应。",
        "3. 「票据号」列：票据的唯一编号。",
        "4. 「票面金额(D)」列：填写面值金额。",
        "5. 「计量日(E)」列：默认为资产负债表日，可修改。",
        "6. 「到期日(F)」列：票据/应收款到期日。",
        "7. 「市场贴现利率(H)」列：使用市场可观察利率（如银行同期贴现利率）。",
        "8. 「公允价值层次(L)」列：选择第二层次（可观察输入值）或第三层次（不可观察输入值）。",
        "",
        "四、自动计算列（灰色背景，无需填写）",
        "G = F - E（剩余天数）",
        "I = D × H × G ÷ 360（贴现利息）",
        "J = D - I（贴现金额/公允价值）",
        "K = J（期末公允价值）",
        "",
        "五、公允价值层次判定",
        "第二层次：使用可观察市场输入值（如银行公布的贴现利率、同业拆借利率）",
        "第三层次：使用不可观察输入值（需说明估值技术和假设）",
        "",
        "六、审计关注",
        "1. 贴现利率选取是否合理，是否有市场可比利率佐证。",
        "2. 公允价值合计与D5-2期末审定合计的差额即为OCI公允价值变动。",
        "3. 差额公式：OCI变动 = D5-2期末审定合计（票面小计） - D5-4公允价值合计。",
    ],
}


def _get_d5_guidance_text(sheet_code: str) -> list[str]:
    """获取D5 sheet对应的编制说明文本"""
    if sheet_code in _D5_SHEET_GUIDANCE:
        return _D5_SHEET_GUIDANCE[sheet_code]
    return [
        "编制说明",
        "",
        "1. 按表头列名依次填写各字段数据。",
        "2. 金额列填写数值（正数），无需添加千分位。",
        "3. 日期格式：YYYY-MM-DD。",
        "4. 请勿修改表头（第1行列名），否则导入时会校验失败。",
        "5. 空行将被自动跳过，最多支持500行数据。",
    ]

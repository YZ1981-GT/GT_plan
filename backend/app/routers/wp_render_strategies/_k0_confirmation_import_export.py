"""K0 管理循环函证 — 导入导出（3 端点 + unreplied-entities）.

端点：
- GET  /api/workpapers/{wp_id}/k0/unreplied-entities?sheet=K0-5|K0-6
- POST /api/workpapers/{wp_id}/k0/export-template?sheet=K0-5|K0-6
- POST /api/workpapers/{wp_id}/k0/export-data?sheet=K0-5|K0-6
- POST /api/workpapers/{wp_id}/k0/import-data?sheet=K0-5|K0-6

K0-5: 其他应收款替代程序 4 区块（期后收款/期末余额/本期发生额/往来对账）
K0-6: 其他应付款替代程序 4 区块（期后付款/期末余额/本期发生额/往来对账）
每区块分 1 个 sheet，首列「被函证单位」用于按公司归组。
"""
from __future__ import annotations

import json
import logging
from typing import Any
from uuid import uuid4

import sqlalchemy as sa
from fastapi import APIRouter, Depends, File, HTTPException, Query, UploadFile
from fastapi.responses import StreamingResponse
from openpyxl import Workbook, load_workbook
from openpyxl.styles import Font
from openpyxl.utils import get_column_letter
from sqlalchemy.ext.asyncio import AsyncSession

from app.core.database import get_db
from app.deps import get_current_user
from app.models.core import User

from ._cycle_import_export_common import safe_float, safe_str, workbook_to_response

logger = logging.getLogger(__name__)

router = APIRouter(tags=["k0-import-export"])

ROW_LIMIT = 500

_SHEET_NAME_MAP: dict[str, str] = {
    "K0-5": "替代程序K0-5",
    "K0-6": "替代程序K0-6",
}

# ─── K0-5 其他应收款替代程序列定义 ───

_VOUCHER_COLS: list[tuple[str, str, bool]] = [
    ("日期", "voucher_date", False),
    ("凭证编号", "voucher_no", False),
    ("业务内容", "business_desc", False),
    ("对方科目", "counter_account", False),
    ("金额", "voucher_amount", True),
]

_K05_BLOCKS: dict[str, dict[str, Any]] = {
    "block1": {
        "sheet_title": "①期后收款检查",
        "columns": [
            ("被函证单位", "_entity", False),
            ("序号", "seq", True),
            *_VOUCHER_COLS,
            ("银行回单日期编号", "bank_receipt_date_no", False),
            ("收款方", "receipt_payee", False),
            ("金额(回单)", "receipt_amount", True),
            ("期后收回比例", "post_receipt_ratio", False),
            ("索引号", "ref_index", False),
            ("是否异常", "is_abnormal", False),
        ],
    },
    "block2": {
        "sheet_title": "②期末余额支持性证据",
        "columns": [
            ("被函证单位", "_entity", False),
            ("序号", "seq", True),
            *_VOUCHER_COLS,
            ("借款/垫款审批单日期编号", "loan_approval_date_no", False),
            ("是否恰当审批", "loan_approved", False),
            ("借据/协议编号", "agreement_no", False),
            ("对方单位", "counterparty", False),
            ("金额(协议)", "agreement_amount", True),
            ("索引号", "ref_index", False),
            ("是否异常", "is_abnormal", False),
        ],
    },
    "block3": {
        "sheet_title": "③本期发生额检查",
        "columns": [
            ("被函证单位", "_entity", False),
            ("序号", "seq", True),
            *_VOUCHER_COLS,
            ("原始单据日期编号", "original_doc_date_no", False),
            ("事由", "doc_reason", False),
            ("审批凭证", "approval_doc", False),
            ("审批人", "approver", False),
            ("金额(审批)", "approval_amount", True),
            ("索引号", "ref_index", False),
            ("是否异常", "is_abnormal", False),
        ],
    },
    "block4": {
        "sheet_title": "④往来对账/协议证据",
        "columns": [
            ("被函证单位", "_entity", False),
            ("序号", "seq", True),
            *_VOUCHER_COLS,
            ("对账单日期", "reconcile_date", False),
            ("对方余额", "other_balance", True),
            ("本方余额", "self_balance", True),
            ("对账差异", "reconcile_diff", True),
            ("往来协议编号", "agreement_ref_no", False),
            ("签订日期", "agreement_sign_date", False),
            ("索引号", "ref_index", False),
            ("是否异常", "is_abnormal", False),
        ],
    },
}

# ─── K0-6 其他应付款替代程序列定义 ───

_K06_BLOCKS: dict[str, dict[str, Any]] = {
    "block1": {
        "sheet_title": "①期后付款检查",
        "columns": [
            ("被函证单位", "_entity", False),
            ("序号", "seq", True),
            *_VOUCHER_COLS,
            ("付款审批单日期编号", "payment_approval_date_no", False),
            ("是否恰当审批", "payment_approved", False),
            ("银行回单日期", "bank_receipt_date", False),
            ("收款方", "receipt_payee", False),
            ("金额(回单)", "receipt_amount", True),
            ("索引号", "ref_index", False),
            ("是否异常", "is_abnormal", False),
        ],
    },
    "block2": {
        "sheet_title": "②期末余额支持性证据",
        "columns": [
            ("被函证单位", "_entity", False),
            ("序号", "seq", True),
            *_VOUCHER_COLS,
            ("借款/收款审批单日期编号", "loan_approval_date_no", False),
            ("是否恰当审批", "loan_approved", False),
            ("借据/协议编号", "agreement_no", False),
            ("对方单位", "counterparty", False),
            ("金额(协议)", "agreement_amount", True),
            ("索引号", "ref_index", False),
            ("是否异常", "is_abnormal", False),
        ],
    },
    "block3": {
        "sheet_title": "③本期发生额检查",
        "columns": [
            ("被函证单位", "_entity", False),
            ("序号", "seq", True),
            *_VOUCHER_COLS,
            ("原始单据日期编号", "original_doc_date_no", False),
            ("事由", "doc_reason", False),
            ("审批凭证", "approval_doc", False),
            ("审批人", "approver", False),
            ("金额(审批)", "approval_amount", True),
            ("索引号", "ref_index", False),
            ("是否异常", "is_abnormal", False),
        ],
    },
    "block4": {
        "sheet_title": "④往来对账/协议证据",
        "columns": [
            ("被函证单位", "_entity", False),
            ("序号", "seq", True),
            *_VOUCHER_COLS,
            ("对账单日期", "reconcile_date", False),
            ("对方余额", "other_balance", True),
            ("本方余额", "self_balance", True),
            ("对账差异", "reconcile_diff", True),
            ("往来协议编号", "agreement_ref_no", False),
            ("签订日期", "agreement_sign_date", False),
            ("索引号", "ref_index", False),
            ("是否异常", "is_abnormal", False),
        ],
    },
}

_BLOCKS_MAP: dict[str, dict[str, dict[str, Any]]] = {
    "K0-5": _K05_BLOCKS,
    "K0-6": _K06_BLOCKS,
}

_GUIDANCE_K05 = [
    "K0-5 其他应收款替代程序 导入说明",
    "",
    "本工作簿含 4 个区块 sheet（期后收款检查/期末余额支持性证据/本期发生额检查/往来对账协议证据）。",
    "「被函证单位」列用于将行归入对应被函证单位；同名单位的行会合并到同一项目下。",
    "期后收回比例、往来对账比例等由前端根据期末余额自动重算。",
]

_GUIDANCE_K06 = [
    "K0-6 其他应付款替代程序 导入说明",
    "",
    "本工作簿含 4 个区块 sheet（期后付款检查/期末余额支持性证据/本期发生额检查/往来对账协议证据）。",
    "「被函证单位」列用于将行归入对应被函证单位；同名单位的行会合并到同一项目下。",
    "期后付款比例、往来对账比例等由前端根据期末余额自动重算。",
]

_GUIDANCE_MAP: dict[str, list[str]] = {
    "K0-5": _GUIDANCE_K05,
    "K0-6": _GUIDANCE_K06,
}

_ACCOUNT_TYPE_MAP: dict[str, str] = {
    "K0-5": "其他应收款",
    "K0-6": "其他应付款",
}


def _validate_sheet(sheet: str) -> None:
    if sheet not in _SHEET_NAME_MAP:
        raise HTTPException(400, f"不支持的sheet: {sheet}。支持: {sorted(_SHEET_NAME_MAP)}")


def _get_blocks(sheet: str) -> dict[str, dict[str, Any]]:
    return _BLOCKS_MAP[sheet]


async def _load_html_data(db: AsyncSession, wp_id: str, sheet_name: str) -> dict:
    result = await db.execute(
        sa.text(
            "SELECT parsed_data FROM working_paper WHERE id = :wp_id AND is_deleted = false"
        ),
        {"wp_id": wp_id},
    )
    row = result.fetchone()
    if not row or not row.parsed_data:
        return {}
    parsed = row.parsed_data
    if isinstance(parsed, str):
        try:
            parsed = json.loads(parsed)
        except (json.JSONDecodeError, TypeError):
            return {}
    html_data = (parsed or {}).get("html_data", {}) if isinstance(parsed, dict) else {}
    sheet_data = html_data.get(sheet_name)
    return sheet_data if isinstance(sheet_data, dict) else {}


async def _save_html_data(db: AsyncSession, wp_id: str, sheet_name: str, payload: dict) -> None:
    result = await db.execute(
        sa.text(
            "SELECT parsed_data FROM working_paper WHERE id = :wp_id AND is_deleted = false"
        ),
        {"wp_id": wp_id},
    )
    row = result.fetchone()
    if not row:
        raise HTTPException(404, f"底稿不存在: {wp_id}")
    parsed = row.parsed_data or {}
    if isinstance(parsed, str):
        try:
            parsed = json.loads(parsed)
        except (json.JSONDecodeError, TypeError):
            parsed = {}
    if not isinstance(parsed, dict):
        parsed = {}
    parsed.setdefault("html_data", {})
    parsed["html_data"][sheet_name] = payload
    await db.execute(
        sa.text(
            "UPDATE working_paper SET parsed_data = CAST(:pd AS jsonb), updated_at = NOW() "
            "WHERE id = :wp_id"
        ),
        {"pd": json.dumps(parsed, ensure_ascii=False), "wp_id": wp_id},
    )
    await db.commit()


def _append_header_sheet(
    ws, columns: list[tuple[str, str, bool]], *, title: str | None = None
) -> int:
    header_row = 1
    headers = [c[0] for c in columns]
    if title:
        ws.append([title])
        ws.merge_cells(start_row=1, start_column=1, end_row=1, end_column=max(len(headers), 1))
        ws["A1"].font = Font(bold=True, size=12)
        header_row = 2
    ws.append(headers)
    for h in ws[header_row]:
        h.font = Font(bold=True)
    ws.freeze_panes = f"A{header_row + 1}"
    for idx in range(1, len(headers) + 1):
        ws.column_dimensions[get_column_letter(idx)].width = 14
    return header_row


def _add_guidance(wb: Workbook, lines: list[str]) -> None:
    ws = wb.create_sheet("编制说明")
    for line in lines:
        ws.append([line])
    ws.column_dimensions["A"].width = 80


def _build_workbook(sheet: str, companies: list[dict] | None) -> Workbook:
    blocks = _get_blocks(sheet)
    guidance = _GUIDANCE_MAP[sheet]
    wb = Workbook()
    wb.remove(wb.active)  # type: ignore[arg-type]
    for block_key, cfg in blocks.items():
        ws = wb.create_sheet(cfg["sheet_title"])
        columns = cfg["columns"]
        _append_header_sheet(ws, columns)
        for company in companies or []:
            entity = safe_str(company.get("entity_name"))
            for r in company.get(f"{block_key}_rows", []) or []:
                out_row = []
                for _header, key, is_num in columns:
                    if key == "_entity":
                        out_row.append(entity)
                    elif is_num:
                        out_row.append(safe_float(r.get(key)))
                    else:
                        out_row.append(safe_str(r.get(key)))
                ws.append(out_row)
    _add_guidance(wb, guidance)
    return wb


def content_io(content: bytes):
    import io
    return io.BytesIO(content)


def _parse_sheet_rows(
    content: bytes, ws_title: str | None, columns: list[tuple[str, str, bool]], header_row: int
) -> list[dict]:
    wb = load_workbook(content_io(content), read_only=True, data_only=True)
    try:
        ws = wb[ws_title] if ws_title and ws_title in wb.sheetnames else wb.active
        if ws is None:
            return []
        header_cells = next(ws.iter_rows(min_row=header_row, max_row=header_row))
        actual = [str(c.value).strip() if c.value is not None else "" for c in header_cells]
        expected = [c[0] for c in columns]
        missing = [h for h in expected if h not in actual]
        if missing:
            raise ValueError(f"sheet「{ws.title}」缺少列: {', '.join(missing)}")
        rows: list[dict] = []
        for values in ws.iter_rows(min_row=header_row + 1, values_only=True):
            if values is None or all(v is None for v in values):
                continue
            if len(rows) >= ROW_LIMIT:
                break
            vals = list(values) + [None] * max(0, len(actual) - len(values))
            out: dict[str, Any] = {"_row_id": uuid4().hex[:12]}
            for header, key, is_num in columns:
                idx = actual.index(header)
                raw = vals[idx] if idx < len(vals) else None
                out[key] = safe_float(raw) if is_num else safe_str(raw)
            rows.append(out)
        return rows
    finally:
        wb.close()


# ─── 端点 ───


@router.get("/api/workpapers/{wp_id}/k0/unreplied-entities")
async def unreplied_entities(
    wp_id: str,
    sheet: str = Query(...),
    db: AsyncSession = Depends(get_db),
    current_user: User = Depends(get_current_user),
) -> list[dict[str, Any]]:
    """从 K0-1 函证结果汇总表获取未回函被函证单位列表（供 importFromSummary）。"""
    _validate_sheet(sheet)
    # 1. 获取当前底稿所属 project_id
    proj_result = await db.execute(
        sa.text("SELECT project_id FROM working_paper WHERE id = :wp_id AND is_deleted = false"),
        {"wp_id": wp_id},
    )
    proj_row = proj_result.fetchone()
    if not proj_row:
        raise HTTPException(404, f"底稿不存在: {wp_id}")
    project_id = proj_row.project_id

    # 2. 找同项目 K0-1 底稿
    k01_result = await db.execute(
        sa.text("""
            SELECT wp.parsed_data
            FROM working_paper wp
            JOIN wp_index wi ON wi.id = wp.wp_index_id
            WHERE wi.wp_code = 'K0-1'
              AND wp.project_id = :pid
              AND wp.is_deleted = false
              AND wi.is_deleted = false
            LIMIT 1
        """),
        {"pid": str(project_id)},
    )
    k01_row = k01_result.fetchone()
    if not k01_row or not k01_row.parsed_data:
        return []

    # 3. 解析 html_data → confirmation-v1 格式 rows
    parsed = k01_row.parsed_data
    if isinstance(parsed, str):
        try:
            parsed = json.loads(parsed)
        except (json.JSONDecodeError, TypeError):
            return []
    if not isinstance(parsed, dict):
        return []

    html_data = parsed.get("html_data", {})
    rows: list[dict] = []
    for _sheet_key, sheet_data in html_data.items():
        if isinstance(sheet_data, dict) and sheet_data.get("_format") == "confirmation-v1":
            rows = sheet_data.get("rows", [])
            break

    if not rows:
        return []

    # 4. 过滤未回函项目（match_status === '未回函' 或 is_replied === false）
    account_type = _ACCOUNT_TYPE_MAP.get(sheet, "其他应收款")
    unreplied: list[dict[str, Any]] = []
    for r in rows:
        match_status = r.get("match_status", "")
        is_replied = r.get("is_replied")
        if match_status == "未回函" or (is_replied is False and match_status != "相符"):
            unreplied.append({
                "entity_name": r.get("entity_name", ""),
                "confirm_index": r.get("confirm_index", ""),
                "account_type": account_type,
                "amount": r.get("amount", 0),
            })

    return unreplied


@router.post("/api/workpapers/{wp_id}/k0/export-template")
async def export_template(
    wp_id: str,
    sheet: str = Query(...),
    current_user: User = Depends(get_current_user),
) -> StreamingResponse:
    _validate_sheet(sheet)
    wb = _build_workbook(sheet, None)
    label = "其他应收款" if sheet == "K0-5" else "其他应付款"
    return workbook_to_response(wb, f"{sheet}{label}替代程序_模板.xlsx")


@router.post("/api/workpapers/{wp_id}/k0/export-data")
async def export_data(
    wp_id: str,
    sheet: str = Query(...),
    db: AsyncSession = Depends(get_db),
    current_user: User = Depends(get_current_user),
) -> StreamingResponse:
    _validate_sheet(sheet)
    sheet_name = _SHEET_NAME_MAP[sheet]
    data = await _load_html_data(db, wp_id, sheet_name)
    companies = data.get("companies") if isinstance(data.get("companies"), list) else []
    wb = _build_workbook(sheet, companies)
    label = "其他应收款" if sheet == "K0-5" else "其他应付款"
    return workbook_to_response(wb, f"{sheet}{label}替代程序_数据.xlsx")


@router.post("/api/workpapers/{wp_id}/k0/import-data")
async def import_data(
    wp_id: str,
    sheet: str = Query(...),
    file: UploadFile = File(...),
    db: AsyncSession = Depends(get_db),
    current_user: User = Depends(get_current_user),
) -> dict[str, Any]:
    _validate_sheet(sheet)
    if not file.filename or not file.filename.endswith(".xlsx"):
        raise HTTPException(400, "请上传 .xlsx 格式文件")
    content = await file.read()
    if len(content) > 10 * 1024 * 1024:
        raise HTTPException(400, "文件大小不能超过10MB")

    sheet_name = _SHEET_NAME_MAP[sheet]
    blocks = _get_blocks(sheet)
    fmt = "alternative-k05-v1" if sheet == "K0-5" else "alternative-k06-v1"
    account_type = _ACCOUNT_TYPE_MAP[sheet]
    errors: list[str] = []
    entity_map: dict[str, dict[str, Any]] = {}
    total = 0

    for block_key, cfg in blocks.items():
        try:
            rows = _parse_sheet_rows(content, cfg["sheet_title"], cfg["columns"], header_row=1)
        except ValueError as e:
            errors.append(str(e))
            continue
        except Exception:
            raise HTTPException(400, "无法解析xlsx文件")
        for r in rows:
            entity = safe_str(r.pop("_entity", "")) or "未命名被函证单位"
            company = entity_map.setdefault(
                entity,
                {
                    "_company_id": uuid4().hex[:12],
                    "entity_name": entity,
                    "_source": "import",
                    "sampling": {},
                    "balance": {"item_name": account_type},
                    "block1_rows": [],
                    "block2_rows": [],
                    "block3_rows": [],
                    "block4_rows": [],
                    "conclusion": {},
                },
            )
            company[f"{block_key}_rows"].append(r)
            total += 1

    if errors and not entity_map:
        return {"ok": False, "errors": errors, "imported_count": 0}

    companies = list(entity_map.values())
    for i, c in enumerate(companies, start=1):
        c["seq"] = i
    payload = {"_format": fmt, "companies": companies}
    await _save_html_data(db, wp_id, sheet_name, payload)
    out: dict[str, Any] = {"ok": True, "imported_count": total, "errors": errors}
    return out

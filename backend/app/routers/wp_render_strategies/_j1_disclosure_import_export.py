"""J1 附注披露（上市 五、40 / 国企 八、40）—— 多区块导入导出。

与其它 J1 sheet 的差别：披露页不是单表，而是「3 个五列变动表 + 3 个说明文本域」。
故 workbook 沿用 F2 已验收的范式（`_f2_disclosure_import_export.py`）——
**一区块一 sheet**，文本域集中到「文本说明」sheet：

    编制说明                   —— 使用指引（不参与导入）
    应付职工薪酬 / (1)…列示    —— 汇总表（rows 数组）
    (1)短期薪酬 / (2)…列示     —— 短期薪酬明细（rows 数组）
    (2)设定提存计划 / (3)…列示 —— 设定提存计划明细（rows 数组）
    文本说明                   —— 「标识 | 说明区块 | 内容」三列

形态说明（F2 三形态在 J1 的取舍）：J1 三张表都是**可增删行的五列变动表**，列结构
逐字相同 → 三个区块全为 `rows`（数组整表）形态，列头提到变体级 `headers_of()`，
不在每个区块重复三份。F2 的 `override`（按主键名匹配的覆盖 map）与 `dr`（三来源列）
在 J1 无对应表 ——「从审定表/明细表带入」是前端读 J1-1/J1-2 的联动，不是可覆盖的
map 存储。说明域是**单个 item 存 `{key: text}`**（不同于 F2 的一文本域一 item），
故它不是 `_Block`，由 `_import_texts` / `texts_of()` 单独处理。

存储契约（与前端 `useJ1DisclosureSections.ts` 一致）：
- 全部披露 item 存在 **`remark`** 字段，`conclusion` 为 null
- item_id = `J1-disc-{variant}-{suffix}`，suffix ∈ summary / short-term / post-employment / notes
- 落库字段集 = `_ROW_FIELDS`，与前端 `serialize()` 逐字一致（**不含 `endBalance`**：
  派生列不持久化）
- 三类派生量导入时按底稿同口径**重算**，用户填的值不作权威：
  ① 期末列 =期初 + 增加 − 减少 ② 派生父行三列 = Σ 紧邻「其中：」子行（源模板 SUM 公式）
  ③ 合计行由前端 computed 产出，压根不落库 → 不导出、不导入

两条硬约束：
- **行名可改 → 只能按「行标识」列匹配**（骨架里「其他」在国企短期薪酬 / 设定提存里
  重复出现，按标签匹配会静默覆盖），故首列固定为 `行标识(勿改)`。
- **整表覆盖**：导入即以表内容为准（删掉某行 = 删掉该披露行）；但**整表未填任何金额
  须跳过写库**，防「拿空模板只想导文本」误清既有明细。

Spec: .kiro/specs/j1-disclosure-template-alignment/ Task 16.2
"""

from __future__ import annotations

import io
import json
import logging
from dataclasses import dataclass
from typing import Any, Literal
from uuid import uuid4

import sqlalchemy as sa
from openpyxl import Workbook, load_workbook
from openpyxl.styles import Alignment, Font, PatternFill
from openpyxl.utils import get_column_letter
from sqlalchemy.ext.asyncio import AsyncSession

from ._cycle_import_export_common import ROW_LIMIT, safe_float, safe_str

logger = logging.getLogger(__name__)

Variant = Literal["listed", "soe"]

SHEET_LISTED = "J1-note-listed"
SHEET_SOE = "J1-note-soe"

# sheet 标识 → 变体。挂在 `_j1_import_export` 既有三路由的 sheet/sheet_type 分派上，
# **不新建 router**（沿用 F2 模式）。
DISCLOSURE_SHEET_TYPES: dict[str, Variant] = {SHEET_LISTED: "listed", SHEET_SOE: "soe"}

# 导出文件名用的中文标签（`J1_{label}_模板.xlsx`）
DISCLOSURE_FILE_LABELS: dict[str, str] = {
    SHEET_LISTED: "附注披露(上市公司)",
    SHEET_SOE: "附注披露(国有企业)",
}

_GUIDE_SHEET = "编制说明"
_TEXT_SHEET = "文本说明"
_TEXT_HEADERS = ("标识(勿改)", "说明区块", "内容")

_HEADER_FILL = PatternFill("solid", fgColor="F0EDF5")

NOTES_SUFFIX = "notes"

# 落库字段集（与前端 `serialize()` 逐字一致；顺序即导出 JSON 的键序）
_ROW_FIELDS: tuple[str, ...] = (
    "id",
    "label",
    "category",
    "indent",
    "beginBalance",
    "increase",
    "decrease",
)

# ════════════════════════════════════════════════════════════════════
# 列头 / 骨架行 / 说明域 镜像常量
#
# 🔴 双真源风险：以下常量镜像前端源码，靠 `backend/tests/test_j1_disclosure_import_export.py`
# 的契约测试（直接正则读 .vue / .ts 源码逐条比对）守住漂移。改前端务必同步此处。
#   期初/期末列头 → j1/core/J1TabDisclosure{Listed,Soe}.vue 的 begin-label / end-label
#   骨架行 + category → 同两个 .vue 的 DEFAULT_SUMMARY / DEFAULT_SHORT_TERM / DEFAULT_POST
#   说明域(key,title) → composables/j1NoteSectionMap.ts J1_{LISTED,SOE}_NOTE_FIELDS
#   item_id 键        → composables/workpaper/j1/useJ1DisclosureSections.ts
#   落库字段集        → 同上文件的 `serialize()`
# ════════════════════════════════════════════════════════════════════

_KEY_HEADER = "行标识(勿改)"
_LABEL_HEADER = "项目"
_INDENT_HEADER = "其中：子项(1=是)"
_INCREASE_HEADER = "本期增加"
_DECREASE_HEADER = "本期减少"
_DERIVED_SUFFIX = "(公式，不导入)"

# 期初 / 期末列用语按变体分叉（各自源模板口径；附注侧由 `j1MovementColumns()` 投影成
# 「期初余额 / 期末余额」）。一份常量给两变体共用 = 国企侧列头错位（F3 已踩过）。
_PERIOD_LABELS: dict[str, dict[str, str]] = {
    "listed": {"begin": "上年年末数", "end": "期末数"},
    "soe": {"begin": "期初余额", "end": "期末余额"},
}

# (id, label, indent)
_LISTED_SUMMARY_SKELETON: tuple[tuple[str, str, int], ...] = (
    ("s-1", "短期薪酬", 0),
    ("s-2", "离职后福利-设定提存计划", 0),
    ("s-3", "辞退福利", 0),
    ("s-4", "一年内到期的其他福利", 0),
)

_LISTED_SHORT_TERM_SKELETON: tuple[tuple[str, str, int], ...] = (
    ("st-1", "工资、奖金、津贴和补贴", 0),
    ("st-2", "职工福利费", 0),
    ("st-3", "社会保险费", 0),
    ("st-4", "其中：1. 医疗保险费", 1),
    ("st-5", "2. 工伤保险费", 1),
    ("st-6", "3. 生育保险费", 1),
    ("st-7", "住房公积金", 0),
    ("st-8", "工会经费和职工教育经费", 0),
    ("st-9", "短期带薪缺勤", 0),
    ("st-10", "短期利润分享计划", 0),
    ("st-11", "非货币性福利", 0),
    ("st-12", "其他短期薪酬", 0),
)

_LISTED_POST_SKELETON: tuple[tuple[str, str, int], ...] = (
    ("pe-1", "离职后福利", 0),
    ("pe-2", "其中：1. 基本养老保险费", 1),
    ("pe-3", "2. 失业保险费", 1),
    ("pe-4", "3. 企业年金缴费", 1),
    ("pe-5", "4. 其他", 1),
    ("pe-6", "其他长期职工福利（不适用的删除）", 0),
    ("pe-7", "其中：1. xxx", 1),
    ("pe-8", "2. 其他", 1),
)

_SOE_SUMMARY_SKELETON: tuple[tuple[str, str, int], ...] = (
    ("s-1", "短期薪酬", 0),
    ("s-2", "离职后福利-设定提存计划", 0),
    ("s-3", "辞退福利", 0),
    ("s-4", "一年内到期的其他福利", 0),
    ("s-5", "其他", 0),
)

_SOE_SHORT_TERM_SKELETON: tuple[tuple[str, str, int], ...] = (
    ("st-1", "工资、奖金、津贴和补贴", 0),
    ("st-2", "职工福利费", 0),
    ("st-3", "社会保险费", 0),
    ("st-4", "其中：医疗保险费", 1),
    ("st-5", "工伤保险费", 1),
    ("st-6", "生育保险费", 1),
    ("st-7", "其他", 1),
    ("st-8", "住房公积金", 0),
    ("st-9", "工会经费和职工教育经费", 0),
    ("st-10", "短期带薪缺勤", 0),
    ("st-11", "短期利润分享计划", 0),
    ("st-12", "其他短期薪酬", 0),
)

_SOE_POST_SKELETON: tuple[tuple[str, str, int], ...] = (
    ("pe-1", "离职后福利", 0),
    ("pe-2", "其中：基本养老保险费", 1),
    ("pe-3", "失业保险费", 1),
    ("pe-4", "企业年金缴费", 1),
    ("pe-5", "其他", 1),
    ("pe-6", "其他长期职工福利（不适用的删除）", 0),
    ("pe-7", "其中：xxx", 1),
    ("pe-8", "其他", 1),
)

# 说明域：(持久化 key, 文本域标题)。导入优先按 key 匹配，key 被清空时回退按标题匹配。
_LISTED_TEXTS: tuple[tuple[str, str], ...] = (
    ("shortTerm", "短期薪酬说明"),
    ("postEmployment", "设定提存计划说明"),
    ("severance", "辞退福利说明"),
)

_SOE_TEXTS: tuple[tuple[str, str], ...] = (
    ("soeNonMonetary", "非货币性福利说明"),
    ("soeDefinedContribution", "设定提存计划说明"),
    ("soeDefinedBenefit", "设定受益计划说明"),
)


# ════════════════════════════════════════════════════════════════════
# 区块定义（表驱动）
# ════════════════════════════════════════════════════════════════════


@dataclass(frozen=True)
class _Block:
    """一个可导入区块 ↔ 一个 workbook sheet（列结构见 `headers_of(variant)`）。"""

    sheet: str
    """workbook sheet 名（≤31 字符，禁 []:*?/\\）。"""
    suffix: str
    """item_id 后缀（完整 item_id = `J1-disc-{variant}-{suffix}`）。"""
    category: str
    """行的 `category` 字段值（前端按它判断"在选中行后插入"与合计行归属）。"""
    skeleton: tuple[tuple[str, str, int], ...]
    """模板预置骨架行（id, label, indent），逐字镜像底稿 DEFAULT_* 常量。"""
    note: str
    """写入 sheet 首行的提示。"""


_OVERWRITE_NOTE = (
    "本表为**整表覆盖**：导入后本区块以表内容为准（删掉某行即删掉该披露行）。"
    f"请先「导出数据」再改；若整表未填任何金额则本表跳过（不清空原数据）。"
    f"「{_KEY_HEADER}」用于按行匹配（行名可改，故不能按行名匹配），请勿修改；"
    "手工新增行留空即可，系统自动生成。"
)

_PARENT_NOTE = (
    f"「{_INDENT_HEADER}」填 1 表示本行是上一非子项行的「其中：」明细；"
    "父行三列由其下子项之和派生（源模板 SUM 公式），在父行填数不作准。"
)

_SUMMARY_NOTE = (
    "汇总表分类行可改名、可增删。也可在底稿点「从审定表/明细表带入」自动取数"
    "（期初取 J1-1 审定期初，本期增减取 J1-2 审定增减）。"
)

_LISTED_BLOCKS: tuple[_Block, ...] = (
    _Block(
        sheet="应付职工薪酬",
        suffix="summary",
        category="summary",
        skeleton=_LISTED_SUMMARY_SKELETON,
        note=_SUMMARY_NOTE + _OVERWRITE_NOTE,
    ),
    _Block(
        sheet="(1)短期薪酬",
        suffix="short-term",
        category="short_term",
        skeleton=_LISTED_SHORT_TERM_SKELETON,
        note=_PARENT_NOTE + _OVERWRITE_NOTE,
    ),
    _Block(
        sheet="(2)设定提存计划",
        suffix="post-employment",
        category="post_employment",
        skeleton=_LISTED_POST_SKELETON,
        note=_PARENT_NOTE + _OVERWRITE_NOTE,
    ),
)

_SOE_BLOCKS: tuple[_Block, ...] = (
    _Block(
        sheet="(1)应付职工薪酬列示",
        suffix="summary",
        category="summary",
        skeleton=_SOE_SUMMARY_SKELETON,
        note=_SUMMARY_NOTE + _OVERWRITE_NOTE,
    ),
    _Block(
        sheet="(2)短期薪酬列示",
        suffix="short-term",
        category="short_term",
        skeleton=_SOE_SHORT_TERM_SKELETON,
        note=(
            _PARENT_NOTE
            + "国企口径下「非货币性福利」并入「其他短期薪酬」，不单独设行。"
            + _OVERWRITE_NOTE
        ),
    ),
    _Block(
        sheet="(3)设定提存计划列示",
        suffix="post-employment",
        category="post_employment",
        skeleton=_SOE_POST_SKELETON,
        note=_PARENT_NOTE + _OVERWRITE_NOTE,
    ),
)


def blocks_of(variant: Variant) -> tuple[_Block, ...]:
    return _LISTED_BLOCKS if variant == "listed" else _SOE_BLOCKS


def texts_of(variant: Variant) -> tuple[tuple[str, str], ...]:
    return _LISTED_TEXTS if variant == "listed" else _SOE_TEXTS


def headers_of(variant: Variant) -> tuple[str, ...]:
    """七列表头：行标识 | 项目 | 其中子项 | 期初 | 本期增加 | 本期减少 | 期末(公式)。"""
    labels = _PERIOD_LABELS[variant]
    return (
        _KEY_HEADER,
        _LABEL_HEADER,
        _INDENT_HEADER,
        labels["begin"],
        _INCREASE_HEADER,
        _DECREASE_HEADER,
        f"{labels['end']}{_DERIVED_SUFFIX}",
    )


def _numeric_headers(variant: Variant) -> tuple[str, str, str]:
    """可导入的三个金额列（期末列是派生量，不在其中）。"""
    return (_PERIOD_LABELS[variant]["begin"], _INCREASE_HEADER, _DECREASE_HEADER)


# 金额列表头 → 行字段名
def _numeric_map(variant: Variant) -> dict[str, str]:
    begin_h, inc_h, dec_h = _numeric_headers(variant)
    return {begin_h: "beginBalance", inc_h: "increase", dec_h: "decrease"}


def item_id_of(variant: Variant, suffix: str) -> str:
    return f"J1-disc-{variant}-{suffix}"


def _new_row_id(suffix: str) -> str:
    return f"{suffix}-{uuid4().hex[:12]}"


def _n(v: Any) -> float:
    return safe_float(v)


# ════════════════════════════════════════════════════════════════════
# 派生量重算（与底稿 `j1DisclosureRowModel.ts` 同口径）
# ════════════════════════════════════════════════════════════════════


def _parent_child_spans(rows: list[dict[str, Any]]) -> list[tuple[int, list[int]]]:
    """非缩进行其后紧跟 ≥1 个连续缩进行 → 该行是派生父行。"""
    out: list[tuple[int, list[int]]] = []
    for i, row in enumerate(rows):
        if row.get("indent"):
            continue
        children: list[int] = []
        for j in range(i + 1, len(rows)):
            if not rows[j].get("indent"):
                break
            children.append(j)
        if children:
            out.append((i, children))
    return out


def apply_derivations(rows: list[dict[str, Any]]) -> list[dict[str, Any]]:
    """就地重算派生量：父行 = Σ 子行（源模板 SUM），期末 = 期初 + 增加 − 减少。"""
    for i, children in _parent_child_spans(rows):
        for key in ("beginBalance", "increase", "decrease"):
            rows[i][key] = round(sum(_n(rows[j].get(key)) for j in children), 2)
    for row in rows:
        row["endBalance"] = round(
            _n(row.get("beginBalance")) + _n(row.get("increase")) - _n(row.get("decrease")), 2
        )
    return rows


def serialize_rows(rows: list[dict[str, Any]]) -> str:
    """按前端 `serialize()` 的字段集与键序落库（派生列 `endBalance` 不持久化）。"""
    return json.dumps(
        [{k: row.get(k) for k in _ROW_FIELDS} for row in rows], ensure_ascii=False
    )


# ════════════════════════════════════════════════════════════════════
# 读写 helper
# ════════════════════════════════════════════════════════════════════


async def _load_remarks(db: AsyncSession, wp_id: str, prefix: str) -> dict[str, str]:
    """一次取回该底稿下所有 `prefix%` 的 remark（item_id → remark）。"""
    result = await db.execute(
        sa.text(
            "SELECT item_id, remark FROM checklist_responses "
            "WHERE wp_id = :wp_id AND item_id LIKE :pat"
        ),
        {"wp_id": wp_id, "pat": f"{prefix}%"},
    )
    return {r[0]: (r[1] or "") for r in result.all()}


def _parse_json(raw: str | None, fallback: Any) -> Any:
    if not raw:
        return fallback
    try:
        parsed = json.loads(raw)
    except (json.JSONDecodeError, TypeError):
        return fallback
    return parsed if isinstance(parsed, type(fallback)) else fallback


async def _upsert_remark(
    db: AsyncSession, wp_id: str, project_id: str, item_id: str, remark: str
) -> None:
    await db.execute(
        sa.text("""
            INSERT INTO checklist_responses (id, project_id, wp_id, item_id, remark, updated_at, created_at)
            VALUES (:id, :project_id, :wp_id, :item_id, :remark, NOW(), NOW())
            ON CONFLICT (wp_id, item_id)
            DO UPDATE SET remark = :remark, updated_at = NOW()
        """),
        {
            "id": str(uuid4()),
            "project_id": project_id,
            "wp_id": wp_id,
            "item_id": item_id,
            "remark": remark,
        },
    )


# ════════════════════════════════════════════════════════════════════
# 导出
# ════════════════════════════════════════════════════════════════════


def _style_header(ws: Any, row_idx: int, ncols: int) -> None:
    for c in range(1, ncols + 1):
        cell = ws.cell(row=row_idx, column=c)
        cell.font = Font(bold=True)
        cell.fill = _HEADER_FILL
        cell.alignment = Alignment(horizontal="center", vertical="center", wrap_text=True)


def _autosize(ws: Any, widths: list[int]) -> None:
    for i, w in enumerate(widths, start=1):
        ws.column_dimensions[get_column_letter(i)].width = w


def _note_row(ws: Any, text: str, ncols: int) -> None:
    ws.append([f"提示：{text}"])
    ws.merge_cells(start_row=1, start_column=1, end_row=1, end_column=ncols)
    cell = ws.cell(row=1, column=1)
    cell.font = Font(italic=True, color="8A6D3B")
    cell.alignment = Alignment(wrap_text=True, vertical="top")


def _skeleton_rows(blk: _Block) -> list[list[Any]]:
    """模板行：预置骨架（行标识 / 项目名 / 是否子项），金额列留空。"""
    return [[row_id, label, indent, None, None, None, None] for row_id, label, indent in blk.skeleton]


def _data_rows(payload: Any, blk: _Block) -> list[list[Any]]:
    """已持久化行 → 导出行；合计行不落库，防御性再滤一次。"""
    if not isinstance(payload, list):
        return []
    rows = [r for r in payload if isinstance(r, dict) and not r.get("isSubtotal")]
    if not rows:
        return []
    normalized = apply_derivations(
        [
            {
                "id": safe_str(r.get("id")) or _new_row_id(blk.suffix),
                "label": safe_str(r.get("label")),
                "category": safe_str(r.get("category")) or blk.category,
                "indent": 1 if _n(r.get("indent")) else 0,
                "beginBalance": round(_n(r.get("beginBalance")), 2),
                "increase": round(_n(r.get("increase")), 2),
                "decrease": round(_n(r.get("decrease")), 2),
            }
            for r in rows
        ]
    )
    return [
        [
            r["id"],
            r["label"],
            r["indent"],
            r["beginBalance"],
            r["increase"],
            r["decrease"],
            r["endBalance"],
        ]
        for r in normalized
    ]


def _guidance_lines(variant: Variant) -> list[str]:
    label = "上市公司" if variant == "listed" else "国有企业"
    section = "五、40" if variant == "listed" else "八、40"
    periods = _PERIOD_LABELS[variant]
    lines = [
        f"J1 附注披露信息（{label}）—— 多区块导入导出",
        "",
        "一、每个业务区块对应一个工作表，表名与底稿小节标题一致。",
        "二、导入按**工作表名**定位区块，缺失的工作表跳过（不清空原数据）。",
        f"三、导入即**整表覆盖**该区块：表里没有的行等于被删除；"
        f"但整表未填任何金额时本表跳过（防拿空模板只导文本误清既有明细）。",
        f"四、「{_KEY_HEADER}」列按行匹配（行名允许改，故不能按行名匹配），请勿修改；"
        "手工新增行留空即可，系统自动生成标识。",
        f"五、公式列不采信填写值，一律按底稿口径重算："
        f"「{periods['end']}」= {periods['begin']} + {_INCREASE_HEADER} − {_DECREASE_HEADER}；"
        f"带「其中：」子项的父行三列 = Σ 子行（源模板 SUM 公式）；"
        "合计行由系统生成，不在本表内。",
        f"六、「{_TEXT_SHEET}」按首列标识匹配说明域（标识被清空时回退按「说明区块」名匹配）。",
        f"七、每张表数据行上限 {ROW_LIMIT} 行，超出部分会被截断。",
        f"八、导入后请在底稿点「同步到附注（{section}）」把数据推给附注。",
        "",
        "区块清单：",
    ]
    for blk in blocks_of(variant):
        lines.append(f"  {blk.sheet} — 明细行（整表覆盖），骨架 {len(blk.skeleton)} 行")
    lines.append(f"  {_TEXT_SHEET} — {len(texts_of(variant))} 个说明域")
    return lines


def _write_guide(wb: Workbook, variant: Variant) -> None:
    ws = wb.active
    ws.title = _GUIDE_SHEET
    for line in _guidance_lines(variant):
        ws.append([line])
    ws.cell(row=1, column=1).font = Font(bold=True, size=12)
    ws.column_dimensions["A"].width = 100


def _write_block_sheet(
    wb: Workbook, blk: _Block, variant: Variant, data_rows: list[list[Any]]
) -> None:
    headers = headers_of(variant)
    ws = wb.create_sheet(blk.sheet)
    _note_row(ws, blk.note, len(headers))
    ws.append(list(headers))
    header_row = ws.max_row
    _style_header(ws, header_row, len(headers))
    for r in data_rows:
        ws.append(r)
    ws.freeze_panes = f"A{header_row + 1}"
    _autosize(ws, [18, 34, 14, 16, 16, 16, 18])


def _write_text_sheet(wb: Workbook, variant: Variant, notes_store: dict[str, Any]) -> None:
    ws = wb.create_sheet(_TEXT_SHEET)
    _note_row(
        ws,
        f"按首列「{_TEXT_HEADERS[0]}」匹配说明域（标识被清空时回退按「{_TEXT_HEADERS[1]}」名匹配），"
        "内容列支持多行文本。整表内容全空则跳过（不清空原说明）；"
        "只要有一条填了内容，其余留空即视为清空该说明。",
        len(_TEXT_HEADERS),
    )
    ws.append(list(_TEXT_HEADERS))
    _style_header(ws, ws.max_row, len(_TEXT_HEADERS))
    for key, title in texts_of(variant):
        raw = notes_store.get(key)
        ws.append([key, title, "" if raw is None else str(raw)])
        ws.cell(row=ws.max_row, column=3).alignment = Alignment(wrap_text=True, vertical="top")
    ws.freeze_panes = "A3"
    _autosize(ws, [26, 26, 96])


def build_disclosure_workbook(
    variant: Variant, *, remarks: dict[str, str] | None = None
) -> Workbook:
    """构造多区块 workbook。`remarks=None` 即模板（骨架行 + 空金额）。"""
    wb = Workbook()
    _write_guide(wb, variant)

    for blk in blocks_of(variant):
        data: list[list[Any]] = []
        if remarks is not None:
            data = _data_rows(_parse_json(remarks.get(item_id_of(variant, blk.suffix)), []), blk)
        # 无持久化 / 数据损坏 → 回落骨架，让用户拿到正确的行名与行标识
        _write_block_sheet(wb, blk, variant, data or _skeleton_rows(blk))

    notes_store = (
        _parse_json(remarks.get(item_id_of(variant, NOTES_SUFFIX)), {})
        if remarks is not None
        else {}
    )
    _write_text_sheet(wb, variant, notes_store)
    return wb


async def export_disclosure_data(db: AsyncSession, wp_id: str, variant: Variant) -> Workbook:
    remarks = await _load_remarks(db, wp_id, f"J1-disc-{variant}-")
    return build_disclosure_workbook(variant, remarks=remarks)


# ════════════════════════════════════════════════════════════════════
# 导入
# ════════════════════════════════════════════════════════════════════


def _find_header_row(ws: Any, expected_first: str, max_scan: int = 4) -> int | None:
    """在前 max_scan 行内找表头行（首格 == expected_first）。找不到返回 None。"""
    for r in range(1, min(max_scan, ws.max_row or 1) + 1):
        val = ws.cell(row=r, column=1).value
        if val is not None and str(val).strip() == expected_first:
            return r
    return None


def _header_index(ws: Any, header_row: int, ncols_hint: int) -> dict[str, int]:
    """表头名 → 1-based 列号（按名匹配，容忍列序变化与多余列）。

    🔴 纵向合并的表头（`merge_cells("A1:A2")`）会让 openpyxl **清空 A2** → 读第 2 行
    时该列为空。故本格为空时回退取上一行同列，否则「导入自家导出的模板」必报缺列。
    """
    out: dict[str, int] = {}
    width = max(ws.max_column or ncols_hint, ncols_hint)
    for c in range(1, width + 1):
        raw = ws.cell(row=header_row, column=c).value
        if (raw is None or not str(raw).strip()) and header_row > 1:
            raw = ws.cell(row=header_row - 1, column=c).value
        name = str(raw).strip() if raw is not None else ""
        if name and name not in out:
            out[name] = c
    return out


def _cell(ws: Any, row: int, col: int | None) -> Any:
    return None if not col else ws.cell(row=row, column=col).value


_TRUTHY = {"1", "是", "y", "yes", "true", "√", "其中", "子项"}


def _is_sub_item(val: Any) -> bool:
    if val is None:
        return False
    if isinstance(val, bool):
        return val
    if isinstance(val, (int, float)):
        return bool(val)
    return str(val).strip().lower() in _TRUTHY


def _is_blank(val: Any) -> bool:
    return val is None or (isinstance(val, str) and not val.strip())


def _import_rows(
    ws: Any, blk: _Block, variant: Variant
) -> tuple[list[dict[str, Any]], list[str], int]:
    """**整表覆盖**：返回的行集就是该区块的新内容。

    第三个返回值 `filled` = 填了金额的行数；为 0 时上层跳过写库
    （防「拿空模板只想导文本」误清既有明细）。
    """
    headers = headers_of(variant)
    hdr_row = _find_header_row(ws, headers[0])
    if hdr_row is None:
        return [], [f"「{blk.sheet}」未找到表头「{headers[0]}」，已跳过"], 0
    idx = _header_index(ws, hdr_row, len(headers))
    warnings: list[str] = []
    # 期末列是派生列，缺了无所谓；其余列缺失需提示（按空值导入）
    missing = [h for h in headers[:-1] if h not in idx]
    if missing:
        warnings.append(f"「{blk.sheet}」缺少列 {'、'.join(missing)}，缺列按空值导入")

    numeric = _numeric_map(variant)
    out: list[dict[str, Any]] = []
    seen_ids: set[str] = set()
    filled = 0
    truncated = False
    for r in range(hdr_row + 1, (ws.max_row or hdr_row) + 1):
        if len(out) >= ROW_LIMIT:
            truncated = True
            break
        row_id = safe_str(_cell(ws, r, idx.get(_KEY_HEADER)))
        label = safe_str(_cell(ws, r, idx.get(_LABEL_HEADER)))
        amounts = {field: _cell(ws, r, idx.get(h)) for h, field in numeric.items()}
        has_amount = any(not _is_blank(v) for v in amounts.values())
        if not row_id and not label and not has_amount:
            continue  # 完全空行
        if row_id and row_id in seen_ids:
            warnings.append(f"「{blk.sheet}」行标识「{row_id}」重复，后出现的行已跳过")
            continue
        item: dict[str, Any] = {
            "id": row_id or _new_row_id(blk.suffix),
            "label": label,
            "category": blk.category,
            "indent": 1 if _is_sub_item(_cell(ws, r, idx.get(_INDENT_HEADER))) else 0,
        }
        for field, raw in amounts.items():
            item[field] = round(_n(raw), 2)
        seen_ids.add(item["id"])
        out.append(item)
        if has_amount:
            filled += 1

    if truncated:
        warnings.append(f"「{blk.sheet}」数据行超过 {ROW_LIMIT} 行，已截断")
    apply_derivations(out)
    if filled == 0:
        warnings.append(f"「{blk.sheet}」未填任何金额，已跳过（未改动既有数据）")
    return out, warnings, filled


def _import_texts(ws: Any, variant: Variant) -> tuple[dict[str, str], list[str], int]:
    """说明域：按首列标识匹配，标识被清空时回退按「说明区块」名匹配。

    整表内容全空 → count=0 → 上层跳过写库（不清空既有说明）。
    """
    pairs = texts_of(variant)
    known = {k for k, _t in pairs}
    title_to_key: dict[str, str] = {}
    for k, t in pairs:
        title_to_key.setdefault(t, k)

    hdr_row = _find_header_row(ws, _TEXT_HEADERS[0])
    if hdr_row is None:
        return {}, [f"「{_TEXT_SHEET}」未找到表头「{_TEXT_HEADERS[0]}」，已跳过"], 0
    idx = _header_index(ws, hdr_row, len(_TEXT_HEADERS))
    key_col = idx.get(_TEXT_HEADERS[0], 1)
    title_col = idx.get(_TEXT_HEADERS[1], 2)
    content_col = idx.get(_TEXT_HEADERS[2], 3)

    out: dict[str, str] = {}
    warnings: list[str] = []
    for r in range(hdr_row + 1, (ws.max_row or hdr_row) + 1):
        raw_key = safe_str(_cell(ws, r, key_col))
        raw_title = safe_str(_cell(ws, r, title_col))
        if not raw_key and not raw_title:
            continue
        if raw_key:
            if raw_key not in known:
                warnings.append(f"「{_TEXT_SHEET}」未知标识「{raw_key}」已跳过")
                continue
            key = raw_key
        else:
            fallback = title_to_key.get(raw_title)
            if not fallback:
                warnings.append(f"「{_TEXT_SHEET}」未知区块「{raw_title}」已跳过")
                continue
            key = fallback
        val = _cell(ws, r, content_col)
        out[key] = "" if val is None else str(val)

    if not any(v.strip() for v in out.values()):
        if out:
            warnings.append(f"「{_TEXT_SHEET}」内容全空，已跳过（未改动既有说明）")
        return {}, warnings, 0
    return out, warnings, len(out)


async def import_disclosure_data(
    db: AsyncSession, wp_id: str, variant: Variant, content: bytes
) -> dict[str, Any]:
    """解析多区块 workbook 并逐区块 UPSERT 到 `remark`。

    缺失的 sheet 跳过（不清空原数据）；未知标识/区块名报 warning 并跳过。
    """
    proj = await db.execute(
        sa.text("SELECT project_id FROM working_paper WHERE id = :wp_id AND is_deleted = false"),
        {"wp_id": wp_id},
    )
    project_id = proj.scalar_one_or_none()
    if not project_id:
        return {"ok": False, "errors": ["底稿不存在"], "imported_count": 0}

    try:
        wb = load_workbook(io.BytesIO(content), data_only=True)
    except Exception:  # noqa: BLE001
        return {"ok": False, "errors": ["无法解析 xlsx 文件"], "imported_count": 0}

    warnings: list[str] = []
    imported_blocks: list[str] = []
    total = 0

    try:
        for blk in blocks_of(variant):
            if blk.sheet not in wb.sheetnames:
                continue
            rows, warns, filled = _import_rows(wb[blk.sheet], blk, variant)
            warnings.extend(warns)
            if not filled:
                continue
            await _upsert_remark(
                db,
                wp_id,
                str(project_id),
                item_id_of(variant, blk.suffix),
                serialize_rows(rows),
            )
            imported_blocks.append(blk.sheet)
            total += len(rows)

        if _TEXT_SHEET in wb.sheetnames:
            texts, warns, cnt = _import_texts(wb[_TEXT_SHEET], variant)
            warnings.extend(warns)
            if cnt:
                # 说明域是单 item 存 `{key: text}` → 先读既有值，只覆盖表内出现的键
                notes_id = item_id_of(variant, NOTES_SUFFIX)
                existing = _parse_json(
                    (await _load_remarks(db, wp_id, notes_id)).get(notes_id), {}
                )
                merged = {
                    k: ("" if existing.get(k) is None else str(existing.get(k)))
                    for k, _t in texts_of(variant)
                }
                merged.update(texts)
                await _upsert_remark(
                    db,
                    wp_id,
                    str(project_id),
                    notes_id,
                    json.dumps(merged, ensure_ascii=False),
                )
                imported_blocks.append(_TEXT_SHEET)
                total += cnt
    finally:
        wb.close()

    if not imported_blocks:
        return {
            "ok": False,
            "errors": ["未识别到任何可导入内容，请使用「导出模板」或「导出数据」下载的文件填写"],
            "imported_count": 0,
            "warnings": warnings,
        }

    await db.commit()
    out: dict[str, Any] = {
        "ok": True,
        "imported_count": total,
        "imported_blocks": imported_blocks,
        "errors": [],
    }
    if warnings:
        out["warnings"] = warnings
    return out

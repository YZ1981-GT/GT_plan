"""L0 债务循环函证 — 导入导出（3 端点 + unreplied-entities）.

端点：
- GET  /api/workpapers/{wp_id}/l0/unreplied-entities?sheet=L0-5
- POST /api/workpapers/{wp_id}/l0/export-template?sheet=L0-5
- POST /api/workpapers/{wp_id}/l0/export-data?sheet=L0-5
- POST /api/workpapers/{wp_id}/l0/import-data?sheet=L0-5

L0-5: 长期应付款/借款替代程序 4 区块
  ①期后付款/还款检查
  ②期末余额支持性证据(借款合同/银行对账单)
  ③本期借款检查
  ④抵质押/担保证据
每区块分 1 个 sheet，首列「被函证单位」用于按公司归组。
"""
from __future__ import annotations

import json
import logging
from typing import Any
from uuid import uuid4

import sqlalchemy as sa
from fastapi import APIRouter, Depends, File, HTTPException, Query, UploadFile
from fastapi.responses import StreamingResponse
from openpyxl import Workbook, load_workbook
from openpyxl.styles import Font
from openpyxl.utils import get_column_letter
from sqlalchemy.ext.asyncio import AsyncSession

from app.core.database import get_db
from app.deps import get_current_user
from app.models.core import User

from ._cycle_import_export_common import safe_float, safe_str, workbook_to_response

logger = logging.getLogger(__name__)

router = APIRouter(tags=["l0-import-export"])

ROW_LIMIT = 500

_SHEET_NAME_MAP: dict[str, str] = {
    "L0-5": "替代程序L0-5",
}

# ─── L0-5 长期应付款/借款替代程序列定义 ───

_VOUCHER_COLS: list[tuple[str, str, bool]] = [
    ("日期", "voucher_date", False),
    ("凭证编号", "voucher_no", False),
    ("业务内容", "business_desc", False),
    ("对方科目", "counter_account", False),
    ("金额", "voucher_amount", True),
]

_L05_BLOCKS: dict[str, dict[str, Any]] = {
    "block1": {
        "sheet_title": "①期后付款还款检查",
        "columns": [
            ("被函证单位", "_entity", False),
            ("序号", "seq", True),
            *_VOUCHER_COLS,
            ("还款审批单日期编号", "repay_approval_date_no", False),
            ("是否恰当审批", "repay_approved", False),
            ("银行回单日期", "bank_receipt_date", False),
            ("收款方", "receipt_payee", False),
            ("还款本金", "repayment_principal", True),
            ("还款利息", "repayment_interest", True),
            ("索引号", "ref_index", False),
            ("是否异常", "is_abnormal", False),
        ],
    },
    "block2": {
        "sheet_title": "②期末余额支持性证据",
        "columns": [
            ("被函证单位", "_entity", False),
            ("序号", "seq", True),
            *_VOUCHER_COLS,
            ("借款合同编号", "contract_no", False),
            ("债权人", "creditor", False),
            ("合同金额", "contract_amount", True),
            ("借款期限", "loan_term", False),
            ("银行对账单日期", "statement_date", False),
            ("账面余额", "book_balance", True),
            ("对账差异", "reconcile_diff", True),
            ("索引号", "ref_index", False),
            ("是否异常", "is_abnormal", False),
        ],
    },
    "block3": {
        "sheet_title": "③本期借款检查",
        "columns": [
            ("被函证单位", "_entity", False),
            ("序号", "seq", True),
            *_VOUCHER_COLS,
            ("借款审批单日期编号", "loan_approval_date_no", False),
            ("是否恰当审批", "loan_approved", False),
            ("到账银行回单日期", "arrival_receipt_date", False),
            ("到账金额", "arrival_amount", True),
            ("借款利率", "loan_rate", False),
            ("索引号", "ref_index", False),
            ("是否异常", "is_abnormal", False),
        ],
    },
    "block4": {
        "sheet_title": "④抵质押担保证据",
        "columns": [
            ("被函证单位", "_entity", False),
            ("序号", "seq", True),
            *_VOUCHER_COLS,
            ("抵质押合同编号", "mortgage_contract_no", False),
            ("抵押物", "mortgage_item", False),
            ("抵押金额", "mortgage_amount", True),
            ("担保合同编号", "guarantee_contract_no", False),
            ("担保方", "guarantor", False),
            ("担保金额", "guarantee_amount", True),
            ("他项权证编号", "title_cert_no", False),
            ("索引号", "ref_index", False),
            ("是否异常", "is_abnormal", False),
        ],
    },
}

_BLOCKS_MAP: dict[str, dict[str, dict[str, Any]]] = {
    "L0-5": _L05_BLOCKS,
}

_GUIDANCE_L05 = [
    "L0-5 长期应付款/借款替代程序 导入说明",
    "",
    "本工作簿含 4 个区块 sheet（期后付款还款检查/期末余额支持性证据/本期借款检查/抵质押担保证据）。",
    "「被函证单位」列用于将行归入对应被函证单位；同名单位的行会合并到同一项目下。",
    "期后还款检查比例、抵质押证据检查比例等由前端根据期末余额自动重算。",
    "对账差异 = 账面余额 - 银行对账单余额（前端自动计算）。",
]

_GUIDANCE_MAP: dict[str, list[str]] = {
    "L0-5": _GUIDANCE_L05,
}

_ACCOUNT_TYPE_MAP: dict[str, str] = {
    "L0-5": "长期应付款/借款",
}


def _validate_sheet(sheet: str) -> None:
    if sheet not in _SHEET_NAME_MAP:
        raise HTTPException(400, f"不支持的sheet: {sheet}。支持: {sorted(_SHEET_NAME_MAP)}")


def _get_blocks(sheet: str) -> dict[str, dict[str, Any]]:
    return _BLOCKS_MAP[sheet]


async def _load_html_data(db: AsyncSession, wp_id: str, sheet_name: str) -> dict:
    result = await db.execute(
        sa.text(
            "SELECT parsed_data FROM working_paper WHERE id = :wp_id AND is_deleted = false"
        ),
        {"wp_id": wp_id},
    )
    row = result.fetchone()
    if not row or not row.parsed_data:
        return {}
    parsed = row.parsed_data
    if isinstance(parsed, str):
        try:
            parsed = json.loads(parsed)
        except (json.JSONDecodeError, TypeError):
            return {}
    html_data = (parsed or {}).get("html_data", {}) if isinstance(parsed, dict) else {}
    sheet_data = html_data.get(sheet_name)
    return sheet_data if isinstance(sheet_data, dict) else {}


async def _save_html_data(db: AsyncSession, wp_id: str, sheet_name: str, payload: dict) -> None:
    result = await db.execute(
        sa.text(
            "SELECT parsed_data FROM working_paper WHERE id = :wp_id AND is_deleted = false"
        ),
        {"wp_id": wp_id},
    )
    row = result.fetchone()
    if not row:
        raise HTTPException(404, f"底稿不存在: {wp_id}")
    parsed = row.parsed_data or {}
    if isinstance(parsed, str):
        try:
            parsed = json.loads(parsed)
        except (json.JSONDecodeError, TypeError):
            parsed = {}
    if not isinstance(parsed, dict):
        parsed = {}
    parsed.setdefault("html_data", {})
    parsed["html_data"][sheet_name] = payload
    await db.execute(
        sa.text(
            "UPDATE working_paper SET parsed_data = CAST(:pd AS jsonb), updated_at = NOW() "
            "WHERE id = :wp_id"
        ),
        {"pd": json.dumps(parsed, ensure_ascii=False), "wp_id": wp_id},
    )
    await db.commit()


def _append_header_sheet(
    ws, columns: list[tuple[str, str, bool]], *, title: str | None = None
) -> int:
    header_row = 1
    headers = [c[0] for c in columns]
    if title:
        ws.append([title])
        ws.merge_cells(start_row=1, start_column=1, end_row=1, end_column=max(len(headers), 1))
        ws["A1"].font = Font(bold=True, size=12)
        header_row = 2
    ws.append(headers)
    for h in ws[header_row]:
        h.font = Font(bold=True)
    ws.freeze_panes = f"A{header_row + 1}"
    for idx in range(1, len(headers) + 1):
        ws.column_dimensions[get_column_letter(idx)].width = 14
    return header_row


def _add_guidance(wb: Workbook, lines: list[str]) -> None:
    ws = wb.create_sheet("编制说明")
    for line in lines:
        ws.append([line])
    ws.column_dimensions["A"].width = 80


def _build_workbook(sheet: str, companies: list[dict] | None) -> Workbook:
    blocks = _get_blocks(sheet)
    guidance = _GUIDANCE_MAP[sheet]
    wb = Workbook()
    wb.remove(wb.active)  # type: ignore[arg-type]
    for block_key, cfg in blocks.items():
        ws = wb.create_sheet(cfg["sheet_title"])
        columns = cfg["columns"]
        _append_header_sheet(ws, columns)
        for company in companies or []:
            entity = safe_str(company.get("entity_name"))
            for r in company.get(f"{block_key}_rows", []) or []:
                out_row = []
                for _header, key, is_num in columns:
                    if key == "_entity":
                        out_row.append(entity)
                    elif is_num:
                        out_row.append(safe_float(r.get(key)))
                    else:
                        out_row.append(safe_str(r.get(key)))
                ws.append(out_row)
    _add_guidance(wb, guidance)
    return wb


def content_io(content: bytes):
    import io
    return io.BytesIO(content)


def _parse_sheet_rows(
    content: bytes, ws_title: str | None, columns: list[tuple[str, str, bool]], header_row: int
) -> list[dict]:
    wb = load_workbook(content_io(content), read_only=True, data_only=True)
    try:
        ws = wb[ws_title] if ws_title and ws_title in wb.sheetnames else wb.active
        if ws is None:
            return []
        header_cells = next(ws.iter_rows(min_row=header_row, max_row=header_row))
        actual = [str(c.value).strip() if c.value is not None else "" for c in header_cells]
        expected = [c[0] for c in columns]
        missing = [h for h in expected if h not in actual]
        if missing:
            raise ValueError(f"sheet「{ws.title}」缺少列: {', '.join(missing)}")
        rows: list[dict] = []
        for values in ws.iter_rows(min_row=header_row + 1, values_only=True):
            if values is None or all(v is None for v in values):
                continue
            if len(rows) >= ROW_LIMIT:
                break
            vals = list(values) + [None] * max(0, len(actual) - len(values))
            out: dict[str, Any] = {"_row_id": uuid4().hex[:12]}
            for header, key, is_num in columns:
                idx = actual.index(header)
                raw = vals[idx] if idx < len(vals) else None
                out[key] = safe_float(raw) if is_num else safe_str(raw)
            rows.append(out)
        return rows
    finally:
        wb.close()


# ─── 端点 ───


@router.get("/api/workpapers/{wp_id}/l0/unreplied-entities")
async def unreplied_entities(
    wp_id: str,
    sheet: str = Query(...),
    db: AsyncSession = Depends(get_db),
    current_user: User = Depends(get_current_user),
) -> list[dict[str, Any]]:
    """从 L0-1 函证结果汇总表获取未回函被函证单位列表（供 importFromSummary）。"""
    _validate_sheet(sheet)
    proj_result = await db.execute(
        sa.text("SELECT project_id FROM working_paper WHERE id = :wp_id AND is_deleted = false"),
        {"wp_id": wp_id},
    )
    proj_row = proj_result.fetchone()
    if not proj_row:
        raise HTTPException(404, f"底稿不存在: {wp_id}")
    project_id = proj_row.project_id

    l01_result = await db.execute(
        sa.text("""
            SELECT wp.parsed_data
            FROM working_paper wp
            JOIN wp_index wi ON wi.id = wp.wp_index_id
            WHERE wi.wp_code = 'L0-1'
              AND wp.project_id = :pid
              AND wp.is_deleted = false
              AND wi.is_deleted = false
            LIMIT 1
        """),
        {"pid": str(project_id)},
    )
    l01_row = l01_result.fetchone()
    if not l01_row or not l01_row.parsed_data:
        return []

    parsed = l01_row.parsed_data
    if isinstance(parsed, str):
        try:
            parsed = json.loads(parsed)
        except (json.JSONDecodeError, TypeError):
            return []
    if not isinstance(parsed, dict):
        return []

    html_data = parsed.get("html_data", {})
    rows: list[dict] = []
    for _sheet_key, sheet_data in html_data.items():
        if isinstance(sheet_data, dict) and sheet_data.get("_format") == "confirmation-v1":
            rows = sheet_data.get("rows", [])
            break

    if not rows:
        return []

    account_type = _ACCOUNT_TYPE_MAP.get(sheet, "长期应付款/借款")
    unreplied: list[dict[str, Any]] = []
    for r in rows:
        match_status = r.get("match_status", "")
        is_replied = r.get("is_replied")
        if match_status == "未回函" or (is_replied is False and match_status != "相符"):
            unreplied.append({
                "entity_name": r.get("entity_name", ""),
                "confirm_index": r.get("confirm_index", ""),
                "account_type": account_type,
                "amount": r.get("amount", 0),
            })

    return unreplied


@router.post("/api/workpapers/{wp_id}/l0/export-template")
async def export_template(
    wp_id: str,
    sheet: str = Query(...),
    current_user: User = Depends(get_current_user),
) -> StreamingResponse:
    _validate_sheet(sheet)
    wb = _build_workbook(sheet, None)
    return workbook_to_response(wb, f"{sheet}长期应付款借款替代程序_模板.xlsx")


@router.post("/api/workpapers/{wp_id}/l0/export-data")
async def export_data(
    wp_id: str,
    sheet: str = Query(...),
    db: AsyncSession = Depends(get_db),
    current_user: User = Depends(get_current_user),
) -> StreamingResponse:
    _validate_sheet(sheet)
    sheet_name = _SHEET_NAME_MAP[sheet]
    data = await _load_html_data(db, wp_id, sheet_name)
    companies = data.get("companies") if isinstance(data.get("companies"), list) else []
    wb = _build_workbook(sheet, companies)
    return workbook_to_response(wb, f"{sheet}长期应付款借款替代程序_数据.xlsx")


@router.post("/api/workpapers/{wp_id}/l0/import-data")
async def import_data(
    wp_id: str,
    sheet: str = Query(...),
    file: UploadFile = File(...),
    db: AsyncSession = Depends(get_db),
    current_user: User = Depends(get_current_user),
) -> dict[str, Any]:
    _validate_sheet(sheet)
    if not file.filename or not file.filename.endswith(".xlsx"):
        raise HTTPException(400, "请上传 .xlsx 格式文件")
    content = await file.read()
    if len(content) > 10 * 1024 * 1024:
        raise HTTPException(400, "文件大小不能超过10MB")

    sheet_name = _SHEET_NAME_MAP[sheet]
    blocks = _get_blocks(sheet)
    fmt = "alternative-l05-v1"
    account_type = _ACCOUNT_TYPE_MAP[sheet]
    errors: list[str] = []
    entity_map: dict[str, dict[str, Any]] = {}
    total = 0

    for block_key, cfg in blocks.items():
        try:
            rows = _parse_sheet_rows(content, cfg["sheet_title"], cfg["columns"], header_row=1)
        except ValueError as e:
            errors.append(str(e))
            continue
        except Exception:
            raise HTTPException(400, "无法解析xlsx文件")
        for r in rows:
            entity = safe_str(r.pop("_entity", "")) or "未命名被函证单位"
            company = entity_map.setdefault(
                entity,
                {
                    "_company_id": uuid4().hex[:12],
                    "entity_name": entity,
                    "_source": "import",
                    "sampling": {},
                    "balance": {"item_name": account_type},
                    "block1_rows": [],
                    "block2_rows": [],
                    "block3_rows": [],
                    "block4_rows": [],
                    "conclusion": {},
                },
            )
            company[f"{block_key}_rows"].append(r)
            total += 1

    if errors and not entity_map:
        return {"ok": False, "errors": errors, "imported_count": 0}

    companies = list(entity_map.values())
    for i, c in enumerate(companies, start=1):
        c["seq"] = i
    payload = {"_format": fmt, "companies": companies}
    await _save_html_data(db, wp_id, sheet_name, payload)
    out: dict[str, Any] = {"ok": True, "imported_count": total, "errors": errors}
    return out

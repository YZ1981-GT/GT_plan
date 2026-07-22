"""H2 在建工程 — 导入导出端点.

H2-2：4区段多 sheet（基本信息 / 账面原值 / 审定原值 / 减值与净值），字段对齐 useH2Detail。
其余 sheet：沿用 cycle 通用导入导出。

POST /h2/export-template | export-data | import-data
"""

from __future__ import annotations

import io
import logging
from typing import Any
from uuid import uuid4

from fastapi import APIRouter, Depends, File, HTTPException, Query, UploadFile
from fastapi.responses import StreamingResponse
from openpyxl import Workbook, load_workbook
from openpyxl.styles import Font
from openpyxl.utils import get_column_letter
from sqlalchemy.ext.asyncio import AsyncSession

from app.core.database import get_db
from app.deps import get_current_user
from app.models.core import User

from ._cycle_import_export_common import (
    ROW_LIMIT,
    build_workbook_template,
    create_cycle_import_export_router,
    export_row_by_keys,
    load_json_rows,
    parse_row_by_headers,
    safe_float,
    safe_str,
    upsert_json_rows,
    workbook_to_response,
)

logger = logging.getLogger(__name__)

# ─── H2-2 四区段（对齐前端 useH2Detail / H2TabDetail）────────────────────────

_H2_2_SEGMENTS: list[tuple[str, list[str], list[str]]] = [
    (
        "基本信息",
        [
            "工程名称", "工程编号", "预算金额", "工程类别", "工程状态",
            "开工日期", "预计竣工", "实际竣工", "累计投入", "资金来源",
            "资本化率(%)", "批准文号", "是否抵押", "转固日期", "转入固定资产类别",
            "施工单位", "合同编号", "监理单位", "备注",
        ],
        [
            "name", "projectCode", "budget", "category", "projectStatus",
            "startDate", "plannedEndDate", "actualEndDate", "accumulatedInput", "fundSource",
            "capRate", "approvalDocNo", "isMortgaged", "transferDate", "transferToH1",
            "contractor", "contractNo", "supervisor", "remark",
        ],
    ),
    (
        "账面原值",
        [
            "工程名称", "工程编号",
            "期初余额", "其中:累计资本化(期初)",
            "增加-材料", "增加-人工", "增加-机械", "增加-利息资本化", "增加-其他", "增加合计",
            "转入固定资产", "其他减少", "其中:资本化转出",
            "期末余额", "其中:累计资本化(期末)",
        ],
        [
            "name", "projectCode",
            "cipBegin", "interestBegin",
            "increaseMaterial", "increaseLabor", "increaseMachinery", "increaseInterest", "increaseOther", "increaseTotal",
            "transferAmount", "decrease", "interestDec",
            "cipEnd", "interestEnd",
        ],
    ),
    (
        "审定原值",
        [
            "工程名称", "工程编号",
            "期初调整", "增加调整", "转固调整", "其他减少调整",
            "利息期初调整", "利息增加调整", "资本化转出调整",
            "审定期初", "审定增加", "审定转固", "审定其他减少", "审定期末", "审定累计资本化",
        ],
        [
            "name", "projectCode",
            "adjustBegin", "increaseAdj", "transferAdj", "decreaseAdj",
            "interestOpenAdj", "interestIncAdj", "interestDecAdj",
            "beginAudited", "increaseAudited", "transferAudited", "decreaseAudited", "endAudited", "interestEndAud",
        ],
    ),
    (
        "减值与净值",
        [
            "工程名称", "工程编号",
            "减值期初", "减值增加", "减值减少", "减值未审期末",
            "减值期初调整", "减值增加调整", "减值减少调整", "减值审定期末",
            "期初净值未审", "期初净值审定", "期末净值未审", "期末净值审定",
            "是否抵押",
        ],
        [
            "name", "projectCode",
            "impairmentBegin", "impairmentIncrease", "impairmentDecrease", "impairmentEnd",
            "impairOpenAdj", "impairIncAdj", "impairDecAdj", "impairEndAud",
            "netBeginUnadj", "netBeginAud", "netEndUnadj", "netEndAud",
            "isMortgaged",
        ],
    ),
]

_H2_2_NUMERIC = {
    "budget", "accumulatedInput", "capRate",
    "cipBegin", "interestBegin",
    "increaseMaterial", "increaseLabor", "increaseMachinery", "increaseInterest", "increaseOther", "increaseTotal",
    "transferAmount", "decrease", "interestDec", "cipEnd", "interestEnd",
    "adjustBegin", "increaseAdj", "transferAdj", "decreaseAdj",
    "interestOpenAdj", "interestIncAdj", "interestDecAdj",
    "beginAudited", "increaseAudited", "transferAudited", "decreaseAudited", "endAudited", "interestEndAud",
    "impairmentBegin", "impairmentIncrease", "impairmentDecrease", "impairmentEnd",
    "impairOpenAdj", "impairIncAdj", "impairDecAdj", "impairEndAud",
    "netBeginUnadj", "netBeginAud", "netEndUnadj", "netEndAud",
}

_H2_2_GUIDANCE = [
    "H2-2 在建工程明细表 — 4区段编制说明（对齐致同 / 参照 H1-2）",
    "",
    "①基本信息：预算/进度/状态/资金来源/资本化率/批准文号/是否抵押(Y/N)/转固类别。",
    "②账面原值：期末=期初+增加合计−转入固定资产−其他减少；利息子列单独勾稽。",
    "③审定原值：审定=未审+期初调整/账项调整；期末审定与 H2-1 勾稽。",
    "④减值与净值：净值=原值−减值（无折旧）；抵押项目须同步附注受限披露。",
    "多 sheet 导入按「工程名称+工程编号」合并；公式列导出为结果值，导入后由前端重算。",
]

# ─── 其余 sheet（单表）───────────────────────────────────────────────────────

_H2_3_HEADERS = [
    "序号", "调整事项说明", "类别", "报表项目", "科目代码", "科目名称",
    "附注项目", "摘要", "借方金额", "贷方金额", "对方科目", "索引", "备注",
]
_H2_3_KEYS = [
    "seq", "description", "category", "reportItem", "accountCode", "accountName",
    "noteItem", "summary", "debitAmount", "creditAmount", "contraAccount", "indexRef", "remark",
]

_H2_5_CIP_HEADERS = [
    "序号", "工程项目名称及设备", "在建原值", "利息资本化", "减值准备", "净值",
    "总预算", "投入占预算%", "预计完工/到货", "开工/供货",
    "达可用状态标准", "是否已达可用/在用", "达可用状态时间",
    "未转固原因", "拟转固%", "期后转固情况", "是否异常", "备注",
]
_H2_5_CIP_KEYS = [
    "seq", "name", "cipOriginal", "capitalizedInterest", "impairment", "cipNet",
    "budget", "budgetRatio", "plannedReadyDate", "startDate",
    "readyCriteria", "readyForUse", "readyDate",
    "notTransferReason", "proposedTransferPct", "postPeriodTransfer", "isAbnormal", "remark",
]

_H2_5_HEADERS = [
    "序号", "工程项目名称及设备", "固定资产原值", "累计折旧", "减值准备", "净值",
    "转固时点", "转固金额", "转入资产类别",
    "竣工验收日期", "决算金额", "达可用状态标准", "试生产日期", "正式投产日期",
    "延迟天数",
    "CAS4条件1-实体建造完成", "CAS4条件2-达到设计要求",
    "CAS4条件3-试运转正常", "CAS4条件4-支出已确定", "CAS4条件5-可使用状态",
    "五条件判定", "H1入账金额", "差异", "是否异常", "备注",
]
_H2_5_KEYS = [
    "seq", "name", "faOriginal", "faAccumDep", "faImpairment", "faNet",
    "transferDate", "transferAmount", "assetCategory",
    "acceptanceDate", "acceptanceAmount", "readyCriteria", "trialProductionDate", "officialProductionDate",
    "delayDays",
    "condition1", "condition2",
    "condition3", "condition4", "condition5",
    "allConditionsMet", "h1Amount", "difference", "isAbnormal", "remark",
]

_H2_7_HEADERS = [
    "序号", "项目", "工程项目总造价", "建筑面积(㎡)", "单方造价",
    "可比价1", "可比价2", "可比价3", "可比价均值", "与可比价差异", "差异率(%)",
    "可比价来源索引", "本期增加额", "本期计入现金流量金额", "勾稽差异", "差异原因/关注事项",
]
_H2_7_KEYS = [
    "seq", "name", "totalCost", "buildingArea", "unitCost",
    "comparable1", "comparable2", "comparable3", "comparableAvg", "unitCostDiff", "unitCostDiffRate",
    "comparableSourceIndex", "periodIncrease", "cashFlowAmount", "cashDiff", "diffReason",
]

_H2_8_HEADERS = [
    "序号", "工程名称", "入账日期", "增加方式", "金额", "供应商/施工方",
    "是否关联方", "关联方名称", "关联方关系",
    "合同编号", "监理/进度(出包)", "领料单(自营)", "发票号(设备)", "验收单(设备)", "付款回单",
    "资本化判断", "入账科目", "审计结论", "备注",
]
_H2_8_KEYS = [
    "seq", "name", "date", "additionMethod", "amount", "supplier",
    "isRelatedParty", "relatedPartyName", "relationship",
    "contractNo", "progressDoc", "materialDoc", "invoiceNo", "acceptanceDoc", "paymentRef",
    "capitalizable", "accountEntry", "auditConclusion", "remark",
]

_H2_9_HEADERS = [
    "序号", "工程项目名称", "减少日期", "凭证编号", "减少类型",
    "转入固定资产金额", "其中利息资本化(转固)", "其他减少金额", "其中利息资本化(其他)",
    "审批单日期编号", "是否恰当审批",
    "验收日期", "验收金额", "工程部盖章", "施工方盖章", "监理方盖章",
    "暂估转固", "是否关联方", "关联方名称", "关联方关系", "处置收入",
    "其他关键证据", "查询号", "是否异常", "审计结论", "备注",
]
_H2_9_KEYS = [
    "seq", "name", "decreaseDate", "voucherNo", "decreaseType",
    "transferToFaAmount", "transferInterestCap", "otherDecreaseAmount", "otherInterestCap",
    "approvalRef", "isApproved",
    "acceptanceDate", "acceptanceAmount", "stampEngineering", "stampContractor", "stampSupervisor",
    "isProvisional", "isRelatedParty", "relatedPartyName", "relationship", "disposalIncome",
    "otherEvidence", "queryNo", "isAbnormal", "auditConclusion", "remark",
]

_H2_13_HEADERS = [
    "序号", "抽盘方向", "在建工程名称", "工程编号", "单价",
    "账面数量", "账面金额", "企业盘点数量", "抽盘数量",
    "进度状况", "是否达可使用状态", "已停工时间", "停工原因", "施工状态",
    "盘点结果", "差异原因", "差异金额", "备注",
]
_H2_13_KEYS = [
    "seq", "direction", "projectName", "projectCode", "unitPrice",
    "bookQty", "bookAmount", "clientCountQty", "sampleQty",
    "progressDesc", "readyForUse", "stopDuration", "stopReason", "constructionStatus",
    "stocktakeResult", "diffReason", "diffAmount", "remark",
]

_H2_17_HEADERS = [
    "序号", "交易类型", "关联单位名称", "关联方关系", "资产/服务类别", "在建工程/工程名称",
    "交易金额", "入账价值", "原值", "减值准备", "净值", "处置损益",
    "本期发生额", "累计发生额", "公允/评估价值", "差异率(%)", "同类总额", "占同类%",
    "交易时间", "定价政策", "是否异常", "审批文件", "独董意见", "结论", "备注", "索引号",
]
_H2_17_KEYS = [
    "seq", "transType", "counterparty", "relationship", "assetCategory", "name",
    "transAmount", "bookValue", "originalCost", "impairment", "netValue", "disposalGain",
    "currentAmount", "cumulativeAmount", "appraisedValue", "priceDiffRate", "categoryTotal", "similarRatio",
    "transDate", "pricingPolicy", "hasAnomaly", "approvalDoc", "independentDirectorOpinion",
    "conclusion", "remark", "indexRef",
]

# 单区段别名（兼容按区段单独导入）
_H2_SPECS: dict[str, dict[str, Any]] = {
    "H2-2-base": {
        "item_id": "H2-2-rows",
        "title": "H2-2 在建工程明细表（基本信息）",
        "headers": _H2_2_SEGMENTS[0][1],
        "field_keys": _H2_2_SEGMENTS[0][2],
        "guidance": _H2_2_GUIDANCE,
    },
    "H2-2-cost-unadj": {
        "item_id": "H2-2-rows",
        "title": "H2-2 在建工程明细表（账面原值）",
        "headers": _H2_2_SEGMENTS[1][1],
        "field_keys": _H2_2_SEGMENTS[1][2],
        "guidance": _H2_2_GUIDANCE,
    },
    "H2-2-cost-aud": {
        "item_id": "H2-2-rows",
        "title": "H2-2 在建工程明细表（审定原值）",
        "headers": _H2_2_SEGMENTS[2][1],
        "field_keys": _H2_2_SEGMENTS[2][2],
        "guidance": _H2_2_GUIDANCE,
    },
    "H2-2-impair": {
        "item_id": "H2-2-rows",
        "title": "H2-2 在建工程明细表（减值与净值）",
        "headers": _H2_2_SEGMENTS[3][1],
        "field_keys": _H2_2_SEGMENTS[3][2],
        "guidance": _H2_2_GUIDANCE,
    },
    # 兼容旧区段名 → 映射到新规格（仅作别名注册，导入仍写 H2-2-rows）
    "H2-2-change": {
        "item_id": "H2-2-rows",
        "title": "H2-2 在建工程明细表（账面原值·兼容旧模板）",
        "headers": _H2_2_SEGMENTS[1][1],
        "field_keys": _H2_2_SEGMENTS[1][2],
        "guidance": ["旧「增减变动」区段已并入「账面原值」。请改用最新四区段模板。"],
    },
    "H2-2-transfer": {
        "item_id": "H2-2-rows",
        "title": "H2-2 在建工程明细表（审定原值·兼容旧模板）",
        "headers": _H2_2_SEGMENTS[2][1],
        "field_keys": _H2_2_SEGMENTS[2][2],
        "guidance": ["旧「竣工结转」金额列已并入「账面/审定原值」。转固日期等见「基本信息」。"],
    },
    "H2-3": {
        "item_id": "H2-3-rows",
        "title": "H2-3 调整分录汇总表",
        "headers": _H2_3_HEADERS,
        "field_keys": _H2_3_KEYS,
        "guidance": [
            "H2-3 调整分录汇总表 编制说明",
            "",
            "对齐 Excel：调整事项说明 / 类别(账项调整·报表调整·其他) / 报表项目 / 科目 / 附注 / 借贷 / 索引 / 备注。",
        ],
    },
    "H2-5-cip": {
        "item_id": "H2-5-cip-rows",
        "title": "H2-5 转固时点检查（表一：重大在建挂账）",
        "headers": _H2_5_CIP_HEADERS,
        "field_keys": _H2_5_CIP_KEYS,
        "guidance": ["H2-5 表一：期末仍挂列在建的重大工程；可从 H2-2 期末余额>0 带入。"],
    },
    "H2-5": {
        "item_id": "H2-5-rows",
        "title": "H2-5 转固时点检查（表二：已转固时点）",
        "headers": _H2_5_HEADERS,
        "field_keys": _H2_5_KEYS,
        "guidance": ["H2-5 表二：本期重大转固时点与 CAS4 五条件；与 H2-2/H1 勾稽。"],
    },
    "H2-7": {
        "item_id": "H2-7-rows",
        "title": "H2-7 工程造价比较分析表",
        "headers": _H2_7_HEADERS,
        "field_keys": _H2_7_KEYS,
        "guidance": ["单方造价=总造价/面积；|差异率|>15%须说明。"],
    },
    "H2-8": {
        "item_id": "H2-8-rows",
        "title": "H2-8 增加检查表",
        "headers": _H2_8_HEADERS,
        "field_keys": _H2_8_KEYS,
        "guidance": ["本期增加细节测试；关联方可供 H2-17 带入。"],
    },
    "H2-9": {
        "item_id": "H2-9-rows",
        "title": "H2-9 减少检查表",
        "headers": _H2_9_HEADERS,
        "field_keys": _H2_9_KEYS,
        "guidance": ["转入固定资产 / 其他减少 + 审批验收证据。"],
    },
    "H2-13": {
        "item_id": "H2-13-rows",
        "title": "H2-13 盘点检查表",
        "headers": _H2_13_HEADERS,
        "field_keys": _H2_13_KEYS,
        "guidance": ["双向抽盘；停工工程关注减值→H2-15。"],
    },
    "H2-17": {
        "item_id": "H2-17-rows",
        "title": "H2-17 关联交易检查表",
        "headers": _H2_17_HEADERS,
        "field_keys": _H2_17_KEYS,
        "guidance": ["合并范围外关联方在建工程购入/出售/工程服务。"],
    },
}


def _build_h2_2_multi_sheet(rows: list[dict], *, template_only: bool = False) -> Workbook:
    wb = Workbook()
    wb.remove(wb.active)
    for seg_name, seg_headers, seg_keys in _H2_2_SEGMENTS:
        ws = wb.create_sheet(title=seg_name)
        ws.append([f"H2-2 在建工程明细表 — {seg_name}"])
        ws.merge_cells(start_row=1, start_column=1, end_row=1, end_column=max(len(seg_headers), 1))
        ws["A1"].font = Font(bold=True, size=12)
        ws.append(seg_headers)
        ws.freeze_panes = "A3"
        for col_idx in range(1, len(seg_headers) + 1):
            ws.column_dimensions[get_column_letter(col_idx)].width = 14
        if not template_only:
            for row_data in rows:
                ws.append(export_row_by_keys(row_data, seg_keys))
    ws_guide = wb.create_sheet("编制说明")
    ws_guide.append(["编制说明"])
    ws_guide.append([])
    for line in _H2_2_GUIDANCE:
        ws_guide.append([line])
    ws_guide.column_dimensions["A"].width = 80
    return wb


def _row_merge_key(d: dict) -> str:
    name = safe_str(d.get("name") or d.get("projectName"))
    code = safe_str(d.get("projectCode"))
    return f"{name}||{code}"


def _parse_h2_2_import(content: bytes) -> tuple[list[dict], list[dict]]:
    """解析 H2-2：优先四区段多 sheet，否则按第一区段单表。"""
    wb = load_workbook(io.BytesIO(content), read_only=True, data_only=True)
    errors: list[dict] = []
    merged: dict[str, dict] = {}
    order: list[str] = []

    multi = any(any(seg[0] in s for seg in _H2_2_SEGMENTS) for s in wb.sheetnames)

    if multi:
        for seg_name, seg_headers, seg_keys in _H2_2_SEGMENTS:
            ws = next((wb[s] for s in wb.sheetnames if seg_name in s), None)
            if ws is None:
                errors.append({"row_number": 0, "field": "", "reason": f"缺少工作表: {seg_name}"})
                continue
            for row_idx, row in enumerate(ws.iter_rows(min_row=3, values_only=True), start=3):
                if all(v is None for v in row):
                    continue
                values = list(row) + [None] * max(0, len(seg_headers) - len(row))
                partial: dict[str, Any] = {}
                for col_i, key in enumerate(seg_keys):
                    if col_i >= len(values):
                        break
                    raw = values[col_i]
                    partial[key] = safe_float(raw) if key in _H2_2_NUMERIC else safe_str(raw)
                if not safe_str(partial.get("name")):
                    continue
                key = _row_merge_key(partial)
                if key not in merged:
                    merged[key] = {"rowId": f"row-{uuid4().hex[:10]}"}
                    order.append(key)
                merged[key].update(partial)
    else:
        ws = wb.active
        if ws is None:
            wb.close()
            return [], [{"row_number": 0, "field": "", "reason": "xlsx无活动工作表"}]
        # 尝试用基本信息表头解析
        _, headers, keys = _H2_2_SEGMENTS[0]
        header_row = None
        for i, row in enumerate(ws.iter_rows(min_row=1, max_row=5, values_only=True), start=1):
            cells = [safe_str(c) for c in row if c is not None]
            if "工程名称" in cells:
                header_row = i
                break
        if header_row is None:
            wb.close()
            return [], [{"row_number": 0, "field": "", "reason": "未找到表头「工程名称」"}]
        for row_idx, row in enumerate(ws.iter_rows(min_row=header_row + 1, values_only=True), start=header_row + 1):
            if all(v is None for v in row):
                continue
            parsed = parse_row_by_headers(list(row), headers, keys)
            if not safe_str(parsed.get("name")):
                continue
            for k in list(parsed.keys()):
                if k in _H2_2_NUMERIC:
                    parsed[k] = safe_float(parsed[k])
            key = _row_merge_key(parsed)
            parsed["rowId"] = f"row-{uuid4().hex[:10]}"
            merged[key] = parsed
            order.append(key)

    wb.close()
    result = [merged[k] for k in order]
    if len(result) > ROW_LIMIT:
        result = result[:ROW_LIMIT]
        errors.append({"row_number": ROW_LIMIT, "field": "", "reason": f"超过{ROW_LIMIT}行限制"})
    return result, errors


# ─── Router：H2-2 多 sheet 优先，其余走 cycle ───────────────────────────────

_cycle_router = create_cycle_import_export_router(
    tag="h2-import-export",
    api_prefix="h2",
    specs=_H2_SPECS,
)

router = APIRouter(tags=["h2-import-export"])


@router.post("/api/workpapers/{wp_id}/h2/export-template")
async def h2_export_template(
    wp_id: str,
    sheet: str = Query(...),
    current_user: User = Depends(get_current_user),
) -> StreamingResponse:
    if sheet == "H2-2":
        wb = _build_h2_2_multi_sheet([], template_only=True)
        return workbook_to_response(wb, "H2-2_明细表_模板.xlsx")
    # 委托给 cycle（同路径需手动复用逻辑）
    if sheet not in _H2_SPECS:
        raise HTTPException(400, f"不支持的sheet: {sheet}。支持: H2-2, {sorted(_H2_SPECS)}")
    sp = _H2_SPECS[sheet]
    wb = build_workbook_template(
        sheet,
        sp["headers"],
        title=sp.get("title"),
        guidance=sp.get("guidance"),
    )
    return workbook_to_response(wb, f"{sheet}_模板.xlsx")


@router.post("/api/workpapers/{wp_id}/h2/export-data")
async def h2_export_data(
    wp_id: str,
    sheet: str = Query(...),
    db: AsyncSession = Depends(get_db),
    current_user: User = Depends(get_current_user),
) -> StreamingResponse:
    if sheet == "H2-2":
        rows = await load_json_rows(db, wp_id, "H2-2-rows", field="conclusion")
        if not rows:
            rows = await load_json_rows(db, wp_id, "H2-2-rows", field="remark")
        wb = _build_h2_2_multi_sheet(rows or [], template_only=False)
        return workbook_to_response(wb, "H2-2_明细表_数据.xlsx")
    if sheet not in _H2_SPECS:
        raise HTTPException(400, f"不支持的sheet: {sheet}")
    sp = _H2_SPECS[sheet]
    item_id = sp["item_id"]
    rows = await load_json_rows(db, wp_id, item_id, field="conclusion")
    if not rows:
        rows = await load_json_rows(db, wp_id, item_id, field="remark")
    wb = build_workbook_template(sheet, sp["headers"], title=sp.get("title"), guidance=sp.get("guidance"))
    ws = wb[sheet]
    for d in rows or []:
        ws.append(export_row_by_keys(d, sp["field_keys"]))
    return workbook_to_response(wb, f"{sheet}_数据.xlsx")


@router.post("/api/workpapers/{wp_id}/h2/import-data")
async def h2_import_data(
    wp_id: str,
    sheet: str = Query(...),
    file: UploadFile = File(...),
    db: AsyncSession = Depends(get_db),
    current_user: User = Depends(get_current_user),
) -> dict[str, Any]:
    content = await file.read()
    if sheet == "H2-2":
        rows, errors = _parse_h2_2_import(content)
        await upsert_json_rows(db, wp_id, "H2-2-rows", rows, field="conclusion")
        await upsert_json_rows(db, wp_id, "H2-2-rows", rows, field="remark")
        return {
            "success": True,
            "imported_count": len(rows),
            "warning": "; ".join(e.get("reason", "") for e in errors) if errors else None,
            "errors": errors,
        }
    if sheet not in _H2_SPECS:
        raise HTTPException(400, f"不支持的sheet: {sheet}")
    # 单区段：走 cycle 同逻辑（简化内联）
    sp = _H2_SPECS[sheet]
    wb = load_workbook(io.BytesIO(content), read_only=True, data_only=True)
    ws = wb.active
    if ws is None:
        raise HTTPException(400, "xlsx无工作表")
    headers = sp["headers"]
    keys = sp["field_keys"]
    header_row = 2
    rows: list[dict] = []
    for row in ws.iter_rows(min_row=header_row + 1, values_only=True):
        if all(v is None for v in row):
            continue
        parsed = parse_row_by_headers(list(row), headers, keys)
        if not any(parsed.values()):
            continue
        if "rowId" not in parsed:
            parsed["rowId"] = f"row-{uuid4().hex[:10]}"
        rows.append(parsed)
        if len(rows) >= ROW_LIMIT:
            break
    wb.close()
    item_id = sp["item_id"]
    # 区段导入：与已有 H2-2-rows 按名称合并
    if item_id == "H2-2-rows":
        existing = await load_json_rows(db, wp_id, item_id, field="conclusion") or []
        if not existing:
            existing = await load_json_rows(db, wp_id, item_id, field="remark") or []
        by_key = {_row_merge_key(r): dict(r) for r in existing if _row_merge_key(r)}
        order = list(by_key.keys())
        for r in rows:
            k = _row_merge_key(r)
            if k not in by_key:
                r.setdefault("rowId", f"row-{uuid4().hex[:10]}")
                by_key[k] = r
                order.append(k)
            else:
                by_key[k].update({kk: vv for kk, vv in r.items() if vv not in (None, "")})
        rows = [by_key[k] for k in order]
    await upsert_json_rows(db, wp_id, item_id, rows, field="conclusion")
    await upsert_json_rows(db, wp_id, item_id, rows, field="remark")
    return {"success": True, "imported_count": len(rows)}


# 保留 cycle router 引用，避免 registry 以外重复注册冲突；本文件对外只导出自定义 router
_ = _cycle_router

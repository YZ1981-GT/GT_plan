"""H0 固定资产循环函证 — 导入导出（3 端点）.

端点：
- POST /api/workpapers/{wp_id}/h0/export-template?sheet=H0-5
- POST /api/workpapers/{wp_id}/h0/export-data?sheet=H0-5
- POST /api/workpapers/{wp_id}/h0/import-data?sheet=H0-5

H0-5：4 区块分 4 个 sheet，首列「被函证单位」用于按被函证单位归组。
列定义与前端 blockColumnConfigsH05.ts 对齐。
"""
from __future__ import annotations

import json
import logging
from typing import Any
from uuid import uuid4

import sqlalchemy as sa
from fastapi import APIRouter, Depends, File, HTTPException, Query, UploadFile
from fastapi.responses import StreamingResponse
from openpyxl import Workbook, load_workbook
from openpyxl.styles import Font
from openpyxl.utils import get_column_letter
from sqlalchemy.ext.asyncio import AsyncSession

from app.core.database import get_db
from app.deps import get_current_user
from app.models.core import User

from ._cycle_import_export_common import safe_float, safe_str, workbook_to_response

logger = logging.getLogger(__name__)

router = APIRouter(tags=["h0-import-export"])

ROW_LIMIT = 500

_SHEET_NAME_MAP: dict[str, str] = {
    "H0-5": "替代程序H0-5",
}

_VOUCHER_COLS: list[tuple[str, str, bool]] = [
    ("日期", "voucher_date", False),
    ("凭证编号", "voucher_no", False),
    ("业务内容", "business_desc", False),
    ("对方科目", "counter_account", False),
    ("金额", "voucher_amount", True),
]

_H05_BLOCKS: dict[str, dict[str, Any]] = {
    "block1": {
        "sheet_title": "①期后验收/权属证据",
        "columns": [
            ("被函证单位", "_entity", False),
            ("序号", "seq", True),
            *_VOUCHER_COLS,
            ("验收单日期/编号", "accept_date_no", False),
            ("资产名称", "asset_name", False),
            ("规格型号", "asset_spec", False),
            ("数量", "asset_qty", True),
            ("权属证书编号", "title_cert_no", False),
            ("权属人", "title_owner", False),
            ("取得日期", "title_date", False),
            ("索引号", "ref_index", False),
            ("是否异常", "is_abnormal", False),
        ],
    },
    "block2": {
        "sheet_title": "②期末余额支持性证据",
        "columns": [
            ("被函证单位", "_entity", False),
            ("序号", "seq", True),
            *_VOUCHER_COLS,
            ("合同日期/编号", "contract_date_no", False),
            ("供应商", "contract_vendor", False),
            ("合同金额", "contract_amount", True),
            ("发票日期/编号", "invoice_date_no", False),
            ("发票金额", "invoice_amount", True),
            ("付款日期", "payment_date", False),
            ("付款金额", "payment_amount", True),
            ("索引号", "ref_index", False),
            ("是否异常", "is_abnormal", False),
        ],
    },
    "block3": {
        "sheet_title": "③本期新增资产检查",
        "columns": [
            ("被函证单位", "_entity", False),
            ("序号", "seq", True),
            *_VOUCHER_COLS,
            ("请购审批日期/编号", "req_date_no", False),
            ("是否恰当审批", "req_approved", False),
            ("到货验收日期", "recv_date", False),
            ("资产名称", "recv_asset_name", False),
            ("数量", "recv_qty", True),
            ("转固日期", "cap_date", False),
            ("原值", "cap_cost", True),
            ("索引号", "ref_index", False),
            ("是否异常", "is_abnormal", False),
        ],
    },
    "block4": {
        "sheet_title": "④抵押担保/融资租赁",
        "columns": [
            ("被函证单位", "_entity", False),
            ("序号", "seq", True),
            *_VOUCHER_COLS,
            ("抵押合同编号", "mortgage_contract", False),
            ("抵押权人", "mortgage_holder", False),
            ("担保金额", "mortgage_amount", True),
            ("融资租赁合同编号", "lease_contract", False),
            ("出租方", "lease_lessor", False),
            ("租赁期", "lease_term", False),
            ("他项权证编号", "warrant_no", False),
            ("索引号", "ref_index", False),
            ("是否异常", "is_abnormal", False),
        ],
    },
}

_H05_GUIDANCE = [
    "H0-5 固定资产循环替代程序 导入说明",
    "",
    "本工作簿含 4 个区块 sheet（期后验收权属/期末余额/本期新增/抵押融资租赁）。",
    "「被函证单位」列用于将行归入对应被函证单位；同名单位的行会合并到同一项目下。",
    "权属证据比例、验收证据比例等由前端根据期末余额自动重算。",
]


def _validate_sheet(sheet: str) -> None:
    if sheet not in _SHEET_NAME_MAP:
        raise HTTPException(400, f"不支持的sheet: {sheet}。支持: {sorted(_SHEET_NAME_MAP)}")


async def _load_html_data(db: AsyncSession, wp_id: str, sheet_name: str) -> dict:
    result = await db.execute(
        sa.text(
            "SELECT parsed_data FROM working_paper WHERE id = :wp_id AND is_deleted = false"
        ),
        {"wp_id": wp_id},
    )
    row = result.fetchone()
    if not row or not row.parsed_data:
        return {}
    parsed = row.parsed_data
    if isinstance(parsed, str):
        try:
            parsed = json.loads(parsed)
        except (json.JSONDecodeError, TypeError):
            return {}
    html_data = (parsed or {}).get("html_data", {}) if isinstance(parsed, dict) else {}
    sheet_data = html_data.get(sheet_name)
    return sheet_data if isinstance(sheet_data, dict) else {}


async def _save_html_data(db: AsyncSession, wp_id: str, sheet_name: str, payload: dict) -> None:
    result = await db.execute(
        sa.text(
            "SELECT parsed_data FROM working_paper WHERE id = :wp_id AND is_deleted = false"
        ),
        {"wp_id": wp_id},
    )
    row = result.fetchone()
    if not row:
        raise HTTPException(404, f"底稿不存在: {wp_id}")
    parsed = row.parsed_data or {}
    if isinstance(parsed, str):
        try:
            parsed = json.loads(parsed)
        except (json.JSONDecodeError, TypeError):
            parsed = {}
    if not isinstance(parsed, dict):
        parsed = {}
    parsed.setdefault("html_data", {})
    parsed["html_data"][sheet_name] = payload
    await db.execute(
        sa.text(
            "UPDATE working_paper SET parsed_data = CAST(:pd AS jsonb), updated_at = NOW() "
            "WHERE id = :wp_id"
        ),
        {"pd": json.dumps(parsed, ensure_ascii=False), "wp_id": wp_id},
    )
    await db.commit()


def _append_header_sheet(
    ws, columns: list[tuple[str, str, bool]], *, title: str | None = None
) -> int:
    header_row = 1
    headers = [c[0] for c in columns]
    if title:
        ws.append([title])
        ws.merge_cells(start_row=1, start_column=1, end_row=1, end_column=max(len(headers), 1))
        ws["A1"].font = Font(bold=True, size=12)
        header_row = 2
    ws.append(headers)
    for h in ws[header_row]:
        h.font = Font(bold=True)
    ws.freeze_panes = f"A{header_row + 1}"
    for idx in range(1, len(headers) + 1):
        ws.column_dimensions[get_column_letter(idx)].width = 14
    return header_row


def _add_guidance(wb: Workbook, lines: list[str]) -> None:
    ws = wb.create_sheet("编制说明")
    for line in lines:
        ws.append([line])
    ws.column_dimensions["A"].width = 80


def _build_h05_workbook(companies: list[dict] | None) -> Workbook:
    wb = Workbook()
    wb.remove(wb.active)
    for block_key, cfg in _H05_BLOCKS.items():
        ws = wb.create_sheet(cfg["sheet_title"])
        columns = cfg["columns"]
        _append_header_sheet(ws, columns)
        for company in companies or []:
            entity = safe_str(company.get("entity_name"))
            for r in company.get(f"{block_key}_rows", []) or []:
                out_row = []
                for _header, key, is_num in columns:
                    if key == "_entity":
                        out_row.append(entity)
                    elif is_num:
                        out_row.append(safe_float(r.get(key)))
                    else:
                        out_row.append(safe_str(r.get(key)))
                ws.append(out_row)
    _add_guidance(wb, _H05_GUIDANCE)
    return wb


def content_io(content: bytes):
    import io

    return io.BytesIO(content)


def _parse_sheet_rows(
    content: bytes, ws_title: str | None, columns: list[tuple[str, str, bool]], header_row: int
) -> list[dict]:
    wb = load_workbook(content_io(content), read_only=True, data_only=True)
    try:
        ws = wb[ws_title] if ws_title and ws_title in wb.sheetnames else wb.active
        if ws is None:
            return []
        header_cells = next(ws.iter_rows(min_row=header_row, max_row=header_row))
        actual = [str(c.value).strip() if c.value is not None else "" for c in header_cells]
        expected = [c[0] for c in columns]
        missing = [h for h in expected if h not in actual]
        if missing:
            raise ValueError(f"sheet「{ws.title}」缺少列: {', '.join(missing)}")
        rows: list[dict] = []
        for values in ws.iter_rows(min_row=header_row + 1, values_only=True):
            if values is None or all(v is None for v in values):
                continue
            if len(rows) >= ROW_LIMIT:
                break
            vals = list(values) + [None] * max(0, len(actual) - len(values))
            out: dict[str, Any] = {"_row_id": uuid4().hex[:12]}
            for header, key, is_num in columns:
                idx = actual.index(header)
                raw = vals[idx] if idx < len(vals) else None
                out[key] = safe_float(raw) if is_num else safe_str(raw)
            rows.append(out)
        return rows
    finally:
        wb.close()


@router.post("/api/workpapers/{wp_id}/h0/export-template")
async def export_template(
    wp_id: str,
    sheet: str = Query(...),
    current_user: User = Depends(get_current_user),
) -> StreamingResponse:
    _validate_sheet(sheet)
    wb = _build_h05_workbook(None)
    return workbook_to_response(wb, "H0-5固定资产循环替代程序_模板.xlsx")


@router.post("/api/workpapers/{wp_id}/h0/export-data")
async def export_data(
    wp_id: str,
    sheet: str = Query(...),
    db: AsyncSession = Depends(get_db),
    current_user: User = Depends(get_current_user),
) -> StreamingResponse:
    _validate_sheet(sheet)
    sheet_name = _SHEET_NAME_MAP[sheet]
    data = await _load_html_data(db, wp_id, sheet_name)
    companies = data.get("companies") if isinstance(data.get("companies"), list) else []
    wb = _build_h05_workbook(companies)
    return workbook_to_response(wb, "H0-5固定资产循环替代程序_数据.xlsx")


@router.post("/api/workpapers/{wp_id}/h0/import-data")
async def import_data(
    wp_id: str,
    sheet: str = Query(...),
    file: UploadFile = File(...),
    db: AsyncSession = Depends(get_db),
    current_user: User = Depends(get_current_user),
) -> dict[str, Any]:
    _validate_sheet(sheet)
    if not file.filename or not file.filename.endswith(".xlsx"):
        raise HTTPException(400, "请上传 .xlsx 格式文件")
    content = await file.read()
    if len(content) > 10 * 1024 * 1024:
        raise HTTPException(400, "文件大小不能超过10MB")

    sheet_name = _SHEET_NAME_MAP[sheet]
    errors: list[str] = []
    entity_map: dict[str, dict[str, Any]] = {}
    total = 0
    for block_key, cfg in _H05_BLOCKS.items():
        try:
            rows = _parse_sheet_rows(content, cfg["sheet_title"], cfg["columns"], header_row=1)
        except ValueError as e:
            errors.append(str(e))
            continue
        except Exception:
            raise HTTPException(400, "无法解析xlsx文件")
        for r in rows:
            entity = safe_str(r.pop("_entity", "")) or "未命名被函证单位"
            company = entity_map.setdefault(
                entity,
                {
                    "_company_id": uuid4().hex[:12],
                    "entity_name": entity,
                    "_source": "import",
                    "sampling": {},
                    "balance": {"item_name": "固定资产"},
                    "block1_rows": [],
                    "block2_rows": [],
                    "block3_rows": [],
                    "block4_rows": [],
                    "conclusion": {},
                },
            )
            company[f"{block_key}_rows"].append(r)
            total += 1

    if errors and not entity_map:
        return {"ok": False, "errors": errors, "imported_count": 0}

    companies = list(entity_map.values())
    for i, c in enumerate(companies, start=1):
        c["seq"] = i
    payload = {"_format": "alternative-h05-v1", "companies": companies}
    await _save_html_data(db, wp_id, sheet_name, payload)
    out: dict[str, Any] = {"ok": True, "imported_count": total, "errors": errors}
    return out

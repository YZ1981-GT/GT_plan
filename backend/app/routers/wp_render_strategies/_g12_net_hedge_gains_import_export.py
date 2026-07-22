"""G12 净敞口套期收益 — 导入导出（G12-2 / G12-3 / G12-4 / G12-5 / G12-6）."""

from __future__ import annotations

from io import BytesIO
from typing import Any
from uuid import uuid4

import sqlalchemy as sa
from openpyxl import load_workbook
from openpyxl.styles import Font
from sqlalchemy.ext.asyncio import AsyncSession

from ._cycle_import_export_common import (
    ROW_LIMIT,
    build_workbook_template,
    create_cycle_import_export_router,
    export_row_by_keys,
    load_json_payload,
    load_json_rows,
    parse_row_by_headers,
    upsert_json_payload,
)

_G12_5_SHEET_TEST = "G12-5"
_G12_5_SHEET_META = "样本与结论"
_G12_5_ROWS_ID = "G12-net-exposure-rows"
_G12_5_META_ID = "G12-net-exposure-meta"
_G12_5_CONCLUSION_ID = "G12-net-exposure-conclusion"
_G12_5_NOTE_ID = "G12-net-exposure-audit-note"

_G12_5_GUIDANCE = [
    "G12-5 风险净敞口检查 编制说明",
    "",
    "净头寸≈头寸1−头寸2；同一项目可按币种分行；填写支持性证据与对应套期工具；可与 G12-2/G12-4 交叉验证。",
    f"Sheet「{_G12_5_SHEET_TEST}」：三、测试明细；Sheet「{_G12_5_SHEET_META}」：一/二/四/五区段文本。",
]

_G12_2_HEADERS = [
    "序号", "项目", "净头寸", "套期工具", "行类型",
    "套期工具累计FV变动", "对应销售部分", "对应采购部分", "FV拆分校验",
    "套期调整摊销", "净敞口套期损益", "索引号", "备注",
]
_G12_2_KEYS = [
    "seq", "item", "netPosition", "hedgingInstrument", "rowKind",
    "instrumentFvCumulative", "salesPortion", "purchasePortion", "fvCheckOk",
    "hedgeAdjAmortization", "netHedgePnl", "indexRef", "remark",
]

_G12_3_HEADERS = [
    "序号", "调整事项说明", "类别", "报表项目", "科目代码", "科目名称",
    "附注项目", "借方调整金额", "贷方调整金额", "索引", "备注",
]
_G12_3_KEYS = [
    "seq", "adjustmentDesc", "category", "fsItem", "accountCode", "accountName",
    "noteItem", "debitAmount", "creditAmount", "indexRef", "remark",
]

_G12_4_HEADERS = [
    "套期关系编号", "套期工具名称", "工具类型", "工具期初FV", "工具期末FV", "工具FV变动",
    "估值方法", "公允价值层次", "估值来源", "被套期项目名称", "项目类型", "项目期初FV",
    "项目期末FV", "项目FV变动", "风险因素", "测试方法", "有效性结论",
]
_G12_4_KEYS = [
    "hedgeRelationId", "instrumentName", "instrumentType", "instrumentOpeningFV", "instrumentClosingFV",
    "instrumentFVChange", "instrumentValuationMethod", "instrumentFVLevel", "instrumentValuationSource",
    "itemName", "itemType", "itemOpeningFV", "itemClosingFV", "itemFVChange", "itemRiskFactor",
    "itemTestMethod", "itemEffectivenessConclusion",
]

_G12_5_HEADERS = [
    "序号", "套期关系编号", "项目", "币种", "头寸1描述", "头寸1金额", "头寸2描述", "头寸2金额",
    "净头寸", "证据类型", "支持性证据", "套期工具", "索引号",
]
_G12_5_KEYS = [
    "seq", "hedgeRelationId", "item", "currency", "position1Desc", "position1Amount", "position2Desc", "position2Amount",
    "netPosition", "evidenceType", "supportingEvidence", "hedgingInstrument", "indexRef",
]
_G12_5_REQUIRED_HEADERS = [
    "序号", "项目", "币种", "头寸1描述", "头寸2描述", "净头寸",
]

_G12_6_HEADERS = [
    "序号", "日期", "凭证编号", "业务内容", "关联套期关系", "对方科目", "借方", "贷方",
    "支持性文件描述", "核对1-原始凭证", "核对2-授权批准", "核对3-账务处理", "核对4-套期指定",
    "核对5-套期会计", "核对6-公允价值", "索引号", "是否异常", "异常说明", "风险等级", "来源", "备注",
]
_G12_6_KEYS = [
    "seq", "voucherDate", "voucherNo", "businessContent", "hedgeRelationId", "counterAccount",
    "debitAmount", "creditAmount", "supportingDocDesc", "check1OriginalComplete", "check2Authorization",
    "check3Accounting", "check4HedgeDesignation", "check5HedgeAccounting", "check6FVValuation",
    "indexNo", "isAbnormal", "abnormalDesc", "riskLevel", "source", "remark",
]


def _g12_5_empty_payload() -> dict[str, Any]:
    return {
        "rows": [],
        "testObjective": "",
        "sampleCriteria": "",
        "auditNote": "",
        "conclusion": "",
    }


async def _load_scalar_field(
    db: AsyncSession,
    wp_id: str,
    item_id: str,
    field: str,
) -> str:
    result = await db.execute(
        sa.text(f"SELECT {field} FROM checklist_responses WHERE wp_id = :wp_id AND item_id = :iid LIMIT 1"),
        {"wp_id": wp_id, "iid": item_id},
    )
    row = result.fetchone()
    if not row:
        return ""
    val = getattr(row, field, None)
    return "" if val is None else str(val)


async def _g12_5_export_loader(db: AsyncSession, wp_id: str) -> dict[str, Any]:
    rows = await load_json_rows(db, wp_id, _G12_5_ROWS_ID, field="remark")
    meta_raw = await load_json_payload(db, wp_id, _G12_5_META_ID, field="remark")
    meta = meta_raw if isinstance(meta_raw, dict) else {}
    return {
        "rows": rows,
        "testObjective": str(meta.get("testObjective") or ""),
        "sampleCriteria": str(meta.get("sampleCriteria") or ""),
        "auditNote": await _load_scalar_field(db, wp_id, _G12_5_NOTE_ID, "remark"),
        "conclusion": await _load_scalar_field(db, wp_id, _G12_5_CONCLUSION_ID, "conclusion"),
    }


def _g12_5_enrich_row(raw: dict[str, Any], seq: int) -> dict[str, Any]:
    return {
        "rowId": raw.get("rowId") or f"g12ne-{uuid4().hex[:10]}",
        "seq": raw.get("seq") or seq,
        "hedgeRelationId": raw.get("hedgeRelationId") or "",
        "item": raw.get("item") or "",
        "currency": raw.get("currency") or "",
        "position1Desc": raw.get("position1Desc") or "",
        "position1Amount": raw.get("position1Amount") or "",
        "position2Desc": raw.get("position2Desc") or "",
        "position2Amount": raw.get("position2Amount") or "",
        "netPosition": raw.get("netPosition") or "",
        "netPositionManual": bool(raw.get("netPositionManual") or raw.get("netPosition")),
        "evidenceType": raw.get("evidenceType") or "",
        "supportingEvidence": raw.get("supportingEvidence") or "",
        "hedgingInstrument": raw.get("hedgingInstrument") or "",
        "indexRef": raw.get("indexRef") or "",
    }


def _build_g12_5_workbook(payload: Any, template_only: bool = False):
    data = payload if isinstance(payload, dict) else _g12_5_empty_payload()
    if isinstance(payload, list):
        data = {**_g12_5_empty_payload(), "rows": payload}

    wb = build_workbook_template(
        _G12_5_SHEET_TEST,
        _G12_5_HEADERS,
        title="G12-5 风险净敞口检查表",
        guidance=_G12_5_GUIDANCE,
    )
    ws = wb[_G12_5_SHEET_TEST]
    if not template_only:
        for d in data.get("rows") or []:
            ws.append(export_row_by_keys(d, _G12_5_KEYS))

    ws_meta = wb.create_sheet(_G12_5_SHEET_META)
    meta_rows = [
        ("字段", "值"),
        ("一、测试目标", data.get("testObjective", "")),
        ("二、样本选取标准与规模", data.get("sampleCriteria", "")),
        ("四、审计说明", data.get("auditNote", "")),
        ("五、审计结论", data.get("conclusion", "")),
    ]
    for row in meta_rows:
        ws_meta.append(list(row))
    ws_meta["A1"].font = Font(bold=True)
    ws_meta["B1"].font = Font(bold=True)
    ws_meta.column_dimensions["A"].width = 28
    ws_meta.column_dimensions["B"].width = 80

    return wb


def _parse_g12_5_import(content: bytes) -> tuple[dict[str, Any], list[str], int]:
    errors: list[str] = []
    wb = load_workbook(BytesIO(content), data_only=True)
    payload = _g12_5_empty_payload()
    count = 0

    sheet_name = _G12_5_SHEET_TEST if _G12_5_SHEET_TEST in wb.sheetnames else wb.sheetnames[0]
    ws = wb[sheet_name]
    header_row = 2 if ws["A1"].value and "G12-5" in str(ws["A1"].value) else 1
    actual = [
        str(c.value).strip() if c.value else ""
        for c in next(ws.iter_rows(min_row=header_row, max_row=header_row))
    ]
    missing = [h for h in _G12_5_REQUIRED_HEADERS if h not in actual]
    if missing:
        raise ValueError(f"缺少列: {', '.join(missing)}")

    rows_out: list[dict[str, Any]] = []
    for i, row in enumerate(ws.iter_rows(min_row=header_row + 1, values_only=True), start=1):
        vals = list(row or [])
        if not any(v is not None and str(v).strip() != "" for v in vals):
            continue
        row_dict = parse_row_by_headers(tuple(vals), actual, _G12_5_KEYS, expected_headers=_G12_5_HEADERS)
        if not str(row_dict.get("item") or "").strip() and not str(row_dict.get("netPosition") or "").strip():
            continue
        rows_out.append(_g12_5_enrich_row(row_dict, i))
        count += 1
        if count >= ROW_LIMIT:
            errors.append(f"数据行超过{ROW_LIMIT}行限制，已截断")
            break
    payload["rows"] = rows_out

    if _G12_5_SHEET_META in wb.sheetnames:
        ws_meta = wb[_G12_5_SHEET_META]
        cmap: dict[str, Any] = {}
        for row in ws_meta.iter_rows(min_row=2, values_only=True):
            if not row or row[0] is None:
                continue
            cmap[str(row[0]).strip()] = row[1] if len(row) > 1 else ""
        mapping = {
            "一、测试目标": "testObjective",
            "二、样本选取标准与规模": "sampleCriteria",
            "四、审计说明": "auditNote",
            "五、审计结论": "conclusion",
        }
        for zh, key in mapping.items():
            if zh in cmap and cmap[zh] is not None:
                payload[key] = str(cmap[zh] or "")

    if count == 0 and not any(payload.get(k) for k in ("testObjective", "sampleCriteria", "auditNote", "conclusion")):
        errors.append("未解析到有效测试行或区段文本")
    return payload, errors, count


async def _upsert_scalar_field(
    db: AsyncSession,
    wp_id: str,
    item_id: str,
    field: str,
    value: str,
) -> None:
    proj = await db.execute(
        sa.text(
            "SELECT project_id FROM working_paper WHERE id = :wp_id AND is_deleted = false"
        ),
        {"wp_id": wp_id},
    )
    project_id = proj.scalar_one_or_none()
    if not project_id:
        raise ValueError(f"working_paper not found: {wp_id}")

    await db.execute(
        sa.text(f"""
            INSERT INTO checklist_responses (id, project_id, wp_id, item_id, {field}, updated_at, created_at)
            VALUES (:id, :project_id, :wp_id, :item_id, :payload, NOW(), NOW())
            ON CONFLICT (wp_id, item_id)
            DO UPDATE SET {field} = :payload, updated_at = NOW()
        """),
        {
            "id": str(uuid4()),
            "project_id": str(project_id),
            "wp_id": wp_id,
            "item_id": item_id,
            "payload": value,
        },
    )
    await db.commit()


async def _g12_5_import_handler(db: AsyncSession, wp_id: str, payload: dict[str, Any]) -> None:
    rows = payload.get("rows") or []
    await upsert_json_payload(db, wp_id, _G12_5_ROWS_ID, rows, field="remark")
    await upsert_json_payload(
        db,
        wp_id,
        _G12_5_META_ID,
        {
            "testObjective": payload.get("testObjective") or "",
            "sampleCriteria": payload.get("sampleCriteria") or "",
        },
        field="remark",
    )
    await _upsert_scalar_field(db, wp_id, _G12_5_NOTE_ID, "remark", payload.get("auditNote") or "")
    await _upsert_scalar_field(db, wp_id, _G12_5_CONCLUSION_ID, "conclusion", payload.get("conclusion") or "")


_G12_SPECS: dict[str, dict[str, Any]] = {
    "G12-2": {
        "item_id": "G12-hedge-detail-rows",
        "title": "G12-2 净敞口套期收益明细表",
        "headers": _G12_2_HEADERS,
        "field_keys": _G12_2_KEYS,
        "guidance": [
            "G12-2 净敞口套期收益明细 编制说明",
            "",
            "FV 拆分校验=销售部分+采购部分=累计FV变动；净敞口套期损益=FV行取销售部分+摊销行。",
            "可与 G12-4 公允价值测试、G12-5 净头寸检查交叉验证。",
        ],
    },
    "G12-3": {
        "item_id": "G12-aje-rows",
        "title": "G12-3 调整分录汇总表",
        "headers": _G12_3_HEADERS,
        "field_keys": _G12_3_KEYS,
        "guidance": [
            "G12-3 调整分录汇总 编制说明",
            "",
            "类别：账项调整(AJE)/报表调整(RJE)/其他；借贷须平衡；6103 净额可同步至 G12-1 审定表。",
        ],
    },
    "G12-4": {
        "item_id": "G12-fv-test-rows",
        "title": "G12-4 公允价值测试表",
        "headers": _G12_4_HEADERS,
        "field_keys": _G12_4_KEYS,
        "guidance": [
            "G12-4 公允价值测试 编制说明",
            "",
            "FV变动=期末-期初；套期工具侧与被套期项目侧通过套期关系编号关联。",
        ],
    },
    "G12-5": {
        "item_id": _G12_5_ROWS_ID,
        "title": "G12-5 风险净敞口检查表",
        "headers": _G12_5_HEADERS,
        "field_keys": _G12_5_KEYS,
        "build_workbook": _build_g12_5_workbook,
        "parse_import": _parse_g12_5_import,
        "export_loader": _g12_5_export_loader,
        "import_handler": _g12_5_import_handler,
        "guidance": _G12_5_GUIDANCE,
    },
    "G12-6": {
        "item_id": "G12-voucher-rows",
        "title": "G12-6 凭证检查表",
        "headers": _G12_6_HEADERS,
        "field_keys": _G12_6_KEYS,
        "guidance": [
            "G12-6 凭证检查 编制说明",
            "",
            "核对列填 ✓/✗/未测；任一✗则是否异常=是；检查比例=已查金额÷本期发生额。",
            "核对5=套期会计(CAS24)；核对6=公允价值；可与 G12-2/G12-4/G12-5 交叉验证。",
        ],
    },
}

router = create_cycle_import_export_router(
    tag="g12-import-export",
    api_prefix="g12",
    specs=_G12_SPECS,
    storage_field="remark",
)

"""G11 附注披露 v2 store ↔ 多工作表 Excel（上市/国企）."""

from __future__ import annotations

import io
from typing import Any

from openpyxl import Workbook, load_workbook
from openpyxl.styles import Font

from ._cycle_import_export_common import safe_float, safe_str

_MAIN_HEADERS = ["行键", "项目", "本期发生额", "上期发生额", "备注"]
_MAIN_KEYS = ["rowKey", "label", "currentAmount", "priorAmount", "remark"]

_TD_HEADERS = ["行键", "项目", "本期发生额", "上期发生额"]
_TD_KEYS = ["rowKey", "label", "currentAmount", "priorAmount"]

_NOTE_HEADERS = ["字段", "内容"]

_LISTED_MAIN_SEED: list[tuple[str, str]] = [
    ("equity_method", "权益法核算的长期股权投资收益"),
    ("dispose_lt_equity", "处置长期股权投资产生的投资收益"),
    ("dispose_hfs_lt", "处置划分为持有待售资产的长期股权投资产生的投资收益"),
    ("trading_hold", "交易性金融资产持有期间的投资收益"),
    ("debt_hold_interest", "债权投资持有期间的利息收入"),
    ("oth_debt_hold_interest", "其他债权投资持有期间的利息收入"),
    ("oei_dividend", "其他权益工具投资的股利收入"),
    ("trading_dispose", "处置交易性金融资产取得的投资收益"),
    ("derivative_dispose", "处置衍生金融资产取得的投资收益"),
    ("debt_dispose", "处置债权投资取得的投资收益"),
    ("control_fv_gain", "取得控制权时，股权按公允价值重新计量产生的利得"),
    ("loss_control_fv_gain", "丧失控制权后，剩余股权按公允价值重新计量产生的利得"),
    ("other", "其他"),
]

_SOE_MAIN_SEED: list[tuple[str, str]] = [
    ("equity_method", "权益法核算的长期股权投资收益"),
    ("dispose_lt_equity", "处置长期股权投资产生的投资收益"),
    ("dispose_hfs_lt", "处置划分为持有待售资产的长期股权投资产生的投资收益"),
    ("trading_hold", "交易性金融资产持有期间的投资收益"),
    ("trading_dispose", "处置交易性金融资产取得的投资收益"),
    ("debt_hold_interest", "债权投资持有期间的利息收益"),
    ("oth_debt_hold_interest", "其他债权投资持有期间的利息收益"),
    ("debt_dispose", "债权投资处置收益"),
    ("oth_debt_dispose", "其他债权投资处置收益"),
    ("onfa_hold", "持有其他非流动金融资产期间取得的投资收益"),
    ("onfa_dispose", "处置其他非流动金融资产取得的投资收益"),
    ("control_fv_gain", "取得控制权时，股权按公允价值重新计量产生的利得"),
    ("loss_control_fv_gain", "丧失控制权后，剩余股权按公允价值重新计量产生的利得"),
    ("oei_dividend", "持有其他权益工具投资期间取得的股利收入"),
    ("derivative_dispose", "处置衍生金融资产取得的投资收益"),
    ("hedge_ineffective", "现金流量套期的无效部分的已实现收益"),
    ("debt_restructuring", "债务重组产生的投资收益"),
    ("other", "其他"),
]

_TD_SEED: list[tuple[str, str]] = [
    ("equity_stock", "交易性权益工具投资——股票投资"),
    ("debt_bond", "交易性债务工具投资——债券投资"),
    ("derivative_non_hedge", "衍生工具——未指定为套期关系的衍生工具"),
    ("derivative_hedge", "指定为有效套期关系的衍生工具"),
    ("other", "其他"),
]

_REPATRIATION_PLACEHOLDER = (
    "注：若投资收益汇回有重大限制的，应予以说明。若不存在此类重大限制，也应做出说明。"
)


def _empty_pair() -> dict[str, float]:
    return {"currentAmount": 0.0, "priorAmount": 0.0}


def default_listed_store() -> dict[str, Any]:
    return {
        "version": 2,
        "rows": [
            {
                "rowKey": k,
                "label": label,
                "currentAmount": 0.0,
                "priorAmount": 0.0,
                "remark": "",
            }
            for k, label in _LISTED_MAIN_SEED
        ],
        "tradingDispose": {k: _empty_pair() for k, _ in _TD_SEED},
    }


def default_soe_store() -> dict[str, Any]:
    return {
        "version": 2,
        "rows": [
            {
                "rowKey": k,
                "label": label,
                "currentAmount": 0.0,
                "priorAmount": 0.0,
                "remark": "",
            }
            for k, label in _SOE_MAIN_SEED
        ],
        "repatriationNote": _REPATRIATION_PLACEHOLDER,
    }


def _normalize_rows(raw_rows: Any, seed: list[tuple[str, str]]) -> list[dict[str, Any]]:
    by_key: dict[str, dict[str, Any]] = {}
    if isinstance(raw_rows, list):
        for r in raw_rows:
            if not isinstance(r, dict):
                continue
            key = safe_str(r.get("rowKey"))
            if not key or key == "total":
                continue
            by_key[key] = {
                "rowKey": key,
                "label": safe_str(r.get("label")) or key,
                "currentAmount": safe_float(r.get("currentAmount")),
                "priorAmount": safe_float(r.get("priorAmount")),
                "remark": safe_str(r.get("remark")),
            }
    out: list[dict[str, Any]] = []
    seen: set[str] = set()
    for k, label in seed:
        hit = by_key.get(k)
        out.append(
            hit
            or {
                "rowKey": k,
                "label": label,
                "currentAmount": 0.0,
                "priorAmount": 0.0,
                "remark": "",
            }
        )
        seen.add(k)
    for k, row in by_key.items():
        if k not in seen:
            out.append(row)
    return out


def normalize_listed(payload: Any) -> dict[str, Any]:
    base = default_listed_store()
    if isinstance(payload, list):
        return {**base, "rows": _normalize_rows(payload, _LISTED_MAIN_SEED)}
    if not isinstance(payload, dict):
        return base
    if payload.get("version") == 2 or "rows" in payload:
        td = dict(base["tradingDispose"])
        for k, v in (payload.get("tradingDispose") or {}).items():
            if isinstance(v, dict):
                td[k] = {
                    "currentAmount": safe_float(v.get("currentAmount")),
                    "priorAmount": safe_float(v.get("priorAmount")),
                }
        return {
            "version": 2,
            "rows": _normalize_rows(payload.get("rows"), _LISTED_MAIN_SEED),
            "tradingDispose": td,
        }
    return base


def normalize_soe(payload: Any) -> dict[str, Any]:
    base = default_soe_store()
    if isinstance(payload, list):
        return {**base, "rows": _normalize_rows(payload, _SOE_MAIN_SEED)}
    if not isinstance(payload, dict):
        return base
    if payload.get("version") == 2 or "rows" in payload:
        note = safe_str(payload.get("repatriationNote")) or _REPATRIATION_PLACEHOLDER
        return {
            "version": 2,
            "rows": _normalize_rows(payload.get("rows"), _SOE_MAIN_SEED),
            "repatriationNote": note,
        }
    return base


def _append_sheet(wb: Workbook, title: str, headers: list[str], rows: list[list[Any]]) -> None:
    ws = wb.create_sheet(title)
    ws.append(headers)
    for r in rows:
        ws.append(r)
    ws.freeze_panes = "A2"
    ws["A1"].font = Font(bold=True)


def _find_header_row(ws) -> tuple[int, list[str]]:
    for i, row in enumerate(ws.iter_rows(min_row=1, max_row=10, values_only=True), start=1):
        cells = [safe_str(c) for c in row if c is not None]
        if any(c in ("行键", "rowKey", "字段", "field") for c in cells):
            return i, [safe_str(c) for c in row]
    first = next(ws.iter_rows(min_row=1, max_row=1, values_only=True), ())
    return 1, [safe_str(c) for c in first]


def _parse_table(ws, expected_headers: list[str], field_keys: list[str]) -> list[dict[str, Any]]:
    if ws is None:
        return []
    header_row, headers = _find_header_row(ws)
    rows: list[dict[str, Any]] = []
    for row in ws.iter_rows(min_row=header_row + 1, values_only=True):
        if not any(c is not None and safe_str(c) for c in row):
            continue
        d = {headers[i] if i < len(headers) else f"c{i}": row[i] if i < len(row) else None for i in range(max(len(headers), len(row)))}
        label = safe_str(d.get("项目") or d.get("label"))
        if label in ("合  计", "合计"):
            continue
        key = safe_str(d.get("行键") or d.get("rowKey"))
        if not key and not label:
            continue
        parsed: dict[str, Any] = {"rowKey": key or f"row_{len(rows) + 1}"}
        for hdr, fk in zip(expected_headers, field_keys):
            val = d.get(hdr)
            if fk in ("rowKey", "label", "remark", "field", "content"):
                parsed[fk] = safe_str(val)
            else:
                parsed[fk] = safe_float(val)
        if not parsed.get("label") and label:
            parsed["label"] = label
        rows.append(parsed)
    return rows


def build_listed_disclosure_workbook(payload: Any, *, template_only: bool = False) -> Workbook:
    store = default_listed_store() if template_only else normalize_listed(payload)
    wb = Workbook()
    wb.remove(wb.active)

    main_rows = [
        [r["rowKey"], r["label"], r["currentAmount"], r["priorAmount"], r.get("remark", "")]
        for r in store["rows"]
    ]
    _append_sheet(wb, "主表", _MAIN_HEADERS, main_rows)

    td_map = store.get("tradingDispose") or {}
    td_label = {k: lab for k, lab in _TD_SEED}
    td_rows = [
        [
            k,
            td_label.get(k, k),
            safe_float((td_map.get(k) or {}).get("currentAmount")),
            safe_float((td_map.get(k) or {}).get("priorAmount")),
        ]
        for k, _ in _TD_SEED
    ]
    _append_sheet(wb, "处置交易性明细", _TD_HEADERS, td_rows)

    guide = [
        ["guidance", "有套期业务时须填「处置交易性明细」，并注意与附注九、3保持一致。"],
        ["guidance2", "主表合计应与 G11-1 审定（6111）勾稽；「其中/减」展开行可不导出。"],
    ]
    _append_sheet(wb, "说明", _NOTE_HEADERS, guide)
    return wb


def build_soe_disclosure_workbook(payload: Any, *, template_only: bool = False) -> Workbook:
    store = default_soe_store() if template_only else normalize_soe(payload)
    wb = Workbook()
    wb.remove(wb.active)

    main_rows = [
        [r["rowKey"], r["label"], r["currentAmount"], r["priorAmount"], r.get("remark", "")]
        for r in store["rows"]
    ]
    _append_sheet(wb, "主表", _MAIN_HEADERS, main_rows)

    note_rows = [
        ["repatriationNote", store.get("repatriationNote") or _REPATRIATION_PLACEHOLDER],
    ]
    _append_sheet(wb, "说明", _NOTE_HEADERS, note_rows)
    return wb


def parse_listed_disclosure_workbook(content: bytes) -> tuple[dict[str, Any], list[str], int]:
    errors: list[str] = []
    try:
        wb = load_workbook(io.BytesIO(content), read_only=True, data_only=True)
    except Exception as e:
        return default_listed_store(), [f"无法解析 xlsx: {e}"], 0

    store = default_listed_store()
    count = 0

    main = _parse_table(wb["主表"] if "主表" in wb.sheetnames else None, _MAIN_HEADERS, _MAIN_KEYS)
    if main:
        store["rows"] = _normalize_rows(main, _LISTED_MAIN_SEED)
        # 保留导入中的额外行
        known = {k for k, _ in _LISTED_MAIN_SEED}
        for r in main:
            if r["rowKey"] not in known:
                store["rows"].append(
                    {
                        "rowKey": r["rowKey"],
                        "label": r.get("label") or r["rowKey"],
                        "currentAmount": safe_float(r.get("currentAmount")),
                        "priorAmount": safe_float(r.get("priorAmount")),
                        "remark": safe_str(r.get("remark")),
                    }
                )
        count += len(main)

    td = _parse_table(
        wb["处置交易性明细"] if "处置交易性明细" in wb.sheetnames else None,
        _TD_HEADERS,
        _TD_KEYS,
    )
    if td:
        td_map = dict(store["tradingDispose"])
        for r in td:
            key = r["rowKey"]
            td_map[key] = {
                "currentAmount": safe_float(r.get("currentAmount")),
                "priorAmount": safe_float(r.get("priorAmount")),
            }
        store["tradingDispose"] = td_map
        count += len(td)

    return store, errors, count


def parse_soe_disclosure_workbook(content: bytes) -> tuple[dict[str, Any], list[str], int]:
    errors: list[str] = []
    try:
        wb = load_workbook(io.BytesIO(content), read_only=True, data_only=True)
    except Exception as e:
        return default_soe_store(), [f"无法解析 xlsx: {e}"], 0

    store = default_soe_store()
    count = 0

    main = _parse_table(wb["主表"] if "主表" in wb.sheetnames else None, _MAIN_HEADERS, _MAIN_KEYS)
    if main:
        store["rows"] = _normalize_rows(main, _SOE_MAIN_SEED)
        known = {k for k, _ in _SOE_MAIN_SEED}
        for r in main:
            if r["rowKey"] not in known:
                store["rows"].append(
                    {
                        "rowKey": r["rowKey"],
                        "label": r.get("label") or r["rowKey"],
                        "currentAmount": safe_float(r.get("currentAmount")),
                        "priorAmount": safe_float(r.get("priorAmount")),
                        "remark": safe_str(r.get("remark")),
                    }
                )
        count += len(main)

    note_ws = wb["说明"] if "说明" in wb.sheetnames else None
    if note_ws is not None:
        notes = _parse_table(note_ws, _NOTE_HEADERS, ["field", "content"])
        for n in notes:
            if n.get("field") == "repatriationNote":
                store["repatriationNote"] = safe_str(n.get("content")) or _REPATRIATION_PLACEHOLDER
                count += 1

    return store, errors, count

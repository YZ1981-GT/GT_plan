"""D3 预收账款 — 导入导出三级端点

4个端点：
- POST /api/workpapers/{wp_id}/d3/export-template?sheet={sheet_code}  空白模板xlsx
- POST /api/workpapers/{wp_id}/d3/export-data?sheet={sheet_code}      数据xlsx
- POST /api/workpapers/{wp_id}/d3/import-data?sheet={sheet_code}      解析xlsx写入
- POST /api/workpapers/{wp_id}/d3/import-aux-balance                  从辅助余额表导入

支持sheets: D3-1, D3-2, D3-3, D3-4-debit, D3-4-credit, D3-5, D3-6, D3-7
"""

from __future__ import annotations

import io
import json
import logging
from typing import Any
from uuid import uuid4

from fastapi import APIRouter, Depends, File, HTTPException, Query, UploadFile
from fastapi.responses import StreamingResponse
from openpyxl import Workbook, load_workbook
from sqlalchemy.ext.asyncio import AsyncSession

from app.core.database import get_db
from app.deps import get_current_user
from app.models.core import User

from ._cycle_import_export_common import (
    aging_export_values,
    build_aging_headers,
    match_import_aging,
    resolve_aging_segments,
    subject_aging_periods,
)

logger = logging.getLogger(__name__)

router = APIRouter(tags=["d3-import-export"])

# D3-2 明细表非账龄基础列（账龄列由项目配置动态派生，Task 12.1）
_D3_2_BASE_HEADERS: list[str] = [
    "对方单位名称", "公司代码", "款项性质", "关联方类型",
    "期初未审余额", "期初账项调整", "期初重分类调整",
    "借方发生", "贷方发生", "被审计单位重分类调整",
    "期末账项调整", "期末重分类调整",
    "是否发函", "期后结转", "备注",
]


def _d3_2_dynamic_headers(segments: list[Any]) -> list[str]:
    """D3-2 动态列头：基础列 + 2N 账龄列（期初/期末审定）。Task 12.1。"""
    return _D3_2_BASE_HEADERS + build_aging_headers(segments, subject_aging_periods("D3"))


def _export_d3_2_row_dynamic(data: dict, segments: list[Any]) -> list:
    """D3-2 动态导出行：基础值 + 嵌套账龄值（按 segments 顺序）。Task 12.1。"""
    base = [
        _safe_str(data.get("customerName")),
        _safe_str(data.get("companyCode")),
        _safe_str(data.get("nature")),
        _safe_str(data.get("relationType")),
        _safe_float(data.get("priorUnadjusted")),
        _safe_float(data.get("priorAdjustment")),
        _safe_float(data.get("priorReclass")),
        _safe_float(data.get("debit")),
        _safe_float(data.get("credit")),
        _safe_float(data.get("entityReclass")),
        _safe_float(data.get("endAje")),
        _safe_float(data.get("endRje")),
        _safe_str(data.get("isConfirmed")),
        _safe_float(data.get("postPeriodSettlement")),
        _safe_str(data.get("remark")),
    ]
    return base + aging_export_values(data, segments, subject_aging_periods("D3"))

# ═══════════════════════════════════════════════════════════════════════════════
# Sheet 配置
# ═══════════════════════════════════════════════════════════════════════════════

_ROW_LIMIT = 500

_SUPPORTED_SHEETS: set[str] = {
    "D3-1", "D3-2", "D3-3", "D3-4-debit", "D3-4-credit", "D3-5", "D3-6", "D3-7",
}

_D3_1_SECTION_LABELS: dict[str, str] = {
    "by-nature": "按性质分类",
    "by-aging": "按账龄分类",
}
_D3_1_SECTION_BY_LABEL: dict[str, str] = {v: k for k, v in _D3_1_SECTION_LABELS.items()}

_D3_1_ROWS: list[tuple[str, str, str]] = [
    ("by-nature", "fixed-asset-sales", "预收销售固定资产款"),
    ("by-nature", "land-use-right", "预收销售土地使用权款"),
    ("by-nature", "contract-invalid", "合同不成立时已收取的对价"),
    ("by-nature", "other", "其他"),
    ("by-aging", "within-1-year", "1年以内"),
    ("by-aging", "1-to-2-years", "1至2年"),
    ("by-aging", "2-to-3-years", "2至3年"),
    ("by-aging", "over-3-years", "3年以上"),
]

_D3_1_FIELDS: list[tuple[str, str, bool]] = [
    ("priorUnadjusted", "期初未审", False),
    ("priorAje", "期初AJE", False),
    ("priorRje", "期初RJE", False),
    ("currentUnadjusted", "期末未审", False),
    ("currentAje", "期末AJE", False),
    ("currentRje", "期末RJE", False),
    ("reasonAnalysis", "原因分析", True),
]

_SHEET_HEADERS: dict[str, list[str]] = {
    "D3-1": [
        "区块", "行键", "项目",
        "期初未审", "期初AJE", "期初RJE",
        "期末未审", "期末AJE", "期末RJE",
        "原因分析",
    ],
    "D3-3": [
        "调整事项说明", "类别", "报表项目", "科目名称", "附注项目",
        "借方调整金额", "贷方调整金额", "索引", "备注",
    ],
    "D3-4-debit": ["行键", "项目", "金额", "来源", "备注"],
    "D3-4-credit": ["行键", "项目", "金额", "来源", "备注"],
    "D3-2": [
        "对方单位名称", "公司代码", "款项性质", "关联方类型",
        "期初未审余额", "期初账项调整", "期初重分类调整",
        "期初审定账龄(1年以下)", "期初审定账龄(1~2年)", "期初审定账龄(2~3年)", "期初审定账龄(3年以上)",
        "借方发生", "贷方发生", "被审计单位重分类调整",
        "期末账项调整", "期末重分类调整",
        "审定账龄(1年以下)", "审定账龄(1~2年)", "审定账龄(2~3年)", "审定账龄(3年以上)",
        "是否发函", "期后结转", "备注",
    ],
    "D3-5": [
        "对方单位名称", "期末余额", "账龄", "经济业务说明",
        "未结转或未偿还的原因", "至审计日结转或偿还金额", "处理结论", "处理计划", "备注",
    ],
    "D3-6": [
        "关联方名称", "关联关系", "期初余额", "借方发生", "贷方发生",
        "期末余额", "发生时间及账龄", "发生原因（款项性质）", "索引号", "备注",
    ],
    "D3-7": [
        "客户名称", "日期", "凭证编号", "业务内容", "对方科目",
        "对方明细科目", "贷方金额", "支持性文件",
        "核对内容1", "核对内容2", "核对内容3", "核对内容4", "核对内容5",
        "索引号", "是否异常", "备注说明",
    ],
}

# item_id 映射（JSON 数组类 sheet）
_SHEET_ITEM_ID: dict[str, str] = {
    "D3-2": "D3-det-rows",
    "D3-3": "D3-aje-rows",
    "D3-4-debit": "D3-ana-debit-rows",
    "D3-4-credit": "D3-ana-credit-rows",
    "D3-5": "D3-lt-rows",
    "D3-6": "D3-rp-rows",
    "D3-7": "D3-vc-post-rows",
}


def _get_headers(sheet_code: str) -> list[str]:
    return _SHEET_HEADERS.get(sheet_code, [])


def _validate_sheet(sheet_code: str) -> None:
    if sheet_code not in _SUPPORTED_SHEETS:
        raise HTTPException(
            400,
            f"不支持的sheet: {sheet_code}。支持: {sorted(_SUPPORTED_SHEETS)}",
        )


# ═══════════════════════════════════════════════════════════════════════════════
# 端点
# ═══════════════════════════════════════════════════════════════════════════════


@router.post("/api/workpapers/{wp_id}/d3/export-template")
async def d3_export_template(
    wp_id: str,
    sheet: str = Query(..., description="Sheet编码如D3-2"),
    db: AsyncSession = Depends(get_db),
    current_user: User = Depends(get_current_user),
) -> StreamingResponse:
    """导出空白模板xlsx（含表头+格式，无数据行）"""
    _validate_sheet(sheet)
    # D3-2 明细表：账龄列头按项目账龄配置动态生成（Task 12.1）
    if sheet == "D3-2":
        segments = await resolve_aging_segments(db, wp_id, "D3")
        headers = _d3_2_dynamic_headers(segments)
    else:
        headers = _get_headers(sheet)

    wb = Workbook()
    ws = wb.active
    ws.title = sheet
    ws.append(headers)
    ws.freeze_panes = "A2"
    for col_idx in range(1, len(headers) + 1):
        ws.column_dimensions[chr(64 + min(col_idx, 26))].width = 16

    if sheet == "D3-1":
        for section, row_key, label in _D3_1_ROWS:
            ws.append([
                _D3_1_SECTION_LABELS[section],
                row_key,
                label,
                *[None] * (len(headers) - 3),
            ])

    buffer = io.BytesIO()
    wb.save(buffer)
    buffer.seek(0)

    filename = f"{sheet}_模板.xlsx"
    from urllib.parse import quote
    encoded_filename = quote(filename)
    return StreamingResponse(
        buffer,
        media_type="application/vnd.openxmlformats-officedocument.spreadsheetml.sheet",
        headers={"Content-Disposition": f"attachment; filename*=UTF-8''{encoded_filename}"},
    )


@router.post("/api/workpapers/{wp_id}/d3/export-data")
async def d3_export_data(
    wp_id: str,
    sheet: str = Query(..., description="Sheet编码如D3-2"),
    db: AsyncSession = Depends(get_db),
    current_user: User = Depends(get_current_user),
) -> StreamingResponse:
    """导出当前数据xlsx"""
    _validate_sheet(sheet)

    import sqlalchemy as sa

    # D3-2 明细表：账龄列头按项目账龄配置动态生成（Task 12.1）
    d3_2_segments: list[Any] = []
    if sheet == "D3-2":
        d3_2_segments = await resolve_aging_segments(db, wp_id, "D3")
        headers = _d3_2_dynamic_headers(d3_2_segments)
    else:
        headers = _get_headers(sheet)

    if sheet == "D3-1":
        return await _stream_d3_1_export(wp_id, db, headers, sheet)

    item_id = _SHEET_ITEM_ID[sheet]
    result = await db.execute(
        sa.text(
            "SELECT remark FROM checklist_responses "
            "WHERE wp_id = :wp_id AND item_id = :item_id LIMIT 1"
        ),
        {"wp_id": wp_id, "item_id": item_id},
    )
    row = result.fetchone()
    rows_data: list[dict] = []
    if row and row.remark:
        try:
            rows_data = json.loads(row.remark)
        except (json.JSONDecodeError, TypeError):
            pass

    wb = Workbook()
    ws = wb.active
    ws.title = sheet
    ws.append(headers)
    ws.freeze_panes = "A2"

    for data_row in rows_data:
        if sheet == "D3-2":
            row_values = _export_d3_2_row_dynamic(data_row, d3_2_segments)
        elif sheet == "D3-3":
            row_values = _export_d3_3_row(data_row)
        elif sheet == "D3-4-debit":
            row_values = _export_d3_4_row(data_row)
        elif sheet == "D3-4-credit":
            row_values = _export_d3_4_row(data_row)
        elif sheet == "D3-5":
            row_values = _export_d3_5_row(data_row)
        elif sheet == "D3-6":
            row_values = _export_d3_6_row(data_row)
        else:  # D3-7
            row_values = _export_d3_7_row(data_row)
        ws.append(row_values)

    buffer = io.BytesIO()
    wb.save(buffer)
    buffer.seek(0)

    filename = f"{sheet}_数据.xlsx"
    from urllib.parse import quote
    encoded_filename = quote(filename)
    return StreamingResponse(
        buffer,
        media_type="application/vnd.openxmlformats-officedocument.spreadsheetml.sheet",
        headers={"Content-Disposition": f"attachment; filename*=UTF-8''{encoded_filename}"},
    )


@router.post("/api/workpapers/{wp_id}/d3/import-data")
async def d3_import_data(
    wp_id: str,
    sheet: str = Query(..., description="Sheet编码如D3-2"),
    file: UploadFile = File(...),
    db: AsyncSession = Depends(get_db),
    current_user: User = Depends(get_current_user),
) -> dict[str, Any]:
    """解析上传xlsx，校验格式，写入 checklist_responses"""
    _validate_sheet(sheet)

    if not file.filename or not file.filename.endswith(".xlsx"):
        raise HTTPException(400, "请上传 .xlsx 格式文件")

    content = await file.read()
    if len(content) > 10 * 1024 * 1024:
        raise HTTPException(400, "文件大小不能超过10MB")

    try:
        wb = load_workbook(io.BytesIO(content), read_only=True, data_only=True)
    except Exception:
        raise HTTPException(400, "无法解析xlsx文件，请确认文件格式正确")

    ws = wb.active
    if ws is None:
        raise HTTPException(400, "xlsx文件中无活动工作表")

    actual_headers = [str(cell.value).strip() if cell.value else "" for cell in next(ws.iter_rows(min_row=1, max_row=1))]

    # D3-2：账龄列按 label 动态匹配当前项目配置，验证仅校验基础列（Task 12.2）
    d3_2_segments: list[Any] = []
    skipped_columns: list[str] = []
    if sheet == "D3-2":
        expected_headers = _D3_2_BASE_HEADERS
        d3_2_segments = await resolve_aging_segments(db, wp_id, "D3")
        _, skipped_columns = match_import_aging(
            lambda _h: None, actual_headers, d3_2_segments, subject_aging_periods("D3"),
        )
    else:
        expected_headers = _get_headers(sheet)

    errors: list[str] = []
    missing_cols = [h for h in expected_headers if h not in actual_headers]
    if missing_cols:
        errors.append(f"缺少列: {', '.join(missing_cols)}")

    if errors:
        return {"ok": False, "errors": errors, "imported_count": 0}

    # 解析数据行
    rows_data: list[dict] = []
    truncated = False
    row_count = 0

    for row in ws.iter_rows(min_row=2, values_only=True):
        if all(v is None for v in row):
            continue
        row_count += 1
        if row_count > _ROW_LIMIT:
            truncated = True
            break

        if sheet == "D3-1":
            row_dict = _parse_d3_1_row(row, actual_headers)
            if row_dict:
                rows_data.append(row_dict)
            continue
        if sheet == "D3-2":
            row_dict = _parse_d3_2_row(row, actual_headers, d3_2_segments)
        elif sheet == "D3-3":
            row_dict = _parse_d3_3_row(row, actual_headers)
        elif sheet == "D3-4-debit":
            row_dict = _parse_d3_4_row(row, actual_headers)
        elif sheet == "D3-4-credit":
            row_dict = _parse_d3_4_row(row, actual_headers)
        elif sheet == "D3-5":
            row_dict = _parse_d3_5_row(row, actual_headers)
        elif sheet == "D3-6":
            row_dict = _parse_d3_6_row(row, actual_headers)
        else:
            row_dict = _parse_d3_7_row(row, actual_headers)

        if not row_dict:
            continue
        rows_data.append(row_dict)

    wb.close()

    # 写入 checklist_responses
    import sqlalchemy as sa

    if sheet == "D3-1":
        field_count = await _persist_d3_1_rows(wp_id, rows_data, db)
        await db.commit()
        result_data: dict[str, Any] = {
            "ok": True,
            "imported_count": len(rows_data),
            "field_count": field_count,
            "errors": errors,
        }
        if truncated:
            result_data["warning"] = f"数据行数超过{_ROW_LIMIT}行限制，已截断"
            result_data["truncated"] = True
        return result_data

    item_id = _SHEET_ITEM_ID[sheet]
    remark_json = json.dumps(rows_data, ensure_ascii=False)

    await db.execute(
        sa.text("""
            INSERT INTO checklist_responses (id, wp_id, item_id, remark, updated_at)
            VALUES (:id, :wp_id, :item_id, :remark, NOW())
            ON CONFLICT (wp_id, item_id)
            DO UPDATE SET remark = :remark, updated_at = NOW()
        """),
        {"id": str(uuid4()), "wp_id": wp_id, "item_id": item_id, "remark": remark_json},
    )
    await db.commit()

    result_data: dict[str, Any] = {"ok": True, "imported_count": len(rows_data), "errors": errors}
    if truncated:
        result_data["warning"] = f"数据行数超过{_ROW_LIMIT}行限制，已截断"
        result_data["truncated"] = True
    if skipped_columns:
        result_data["skipped_columns"] = skipped_columns
        result_data["warnings"] = [
            f"以下账龄列未匹配当前账龄配置，已跳过: {', '.join(skipped_columns)}"
        ]

    return result_data


@router.post("/api/workpapers/{wp_id}/d3/import-aux-balance")
async def d3_import_aux_balance(
    wp_id: str,
    db: AsyncSession = Depends(get_db),
    current_user: User = Depends(get_current_user),
) -> dict[str, Any]:
    """从 tb_aux_balance（科目2203，按客户维度）批量导入到D3-2"""
    import sqlalchemy as sa

    # 获取 project_id
    wp_result = await db.execute(
        sa.text("SELECT project_id FROM working_paper WHERE id = :wp_id"),
        {"wp_id": wp_id},
    )
    wp_row = wp_result.fetchone()
    if not wp_row:
        raise HTTPException(404, "底稿不存在")

    project_id = str(wp_row.project_id)

    # 从 tb_aux_balance 查询科目2203按客户维度聚合
    aux_result = await db.execute(
        sa.text("""
            SELECT aux_name,
                   COALESCE(SUM(CASE WHEN period_type = 'opening' THEN balance ELSE 0 END), 0) AS prior_balance,
                   COALESCE(SUM(CASE WHEN period_type = 'closing' THEN balance ELSE 0 END), 0) AS current_balance
            FROM tb_aux_balance
            WHERE project_id = :pid
              AND account_code = '2203'
              AND is_deleted = false
            GROUP BY aux_name
            ORDER BY aux_name
        """),
        {"pid": project_id},
    )
    aux_rows = aux_result.fetchall()

    if not aux_rows:
        return {"ok": True, "imported_count": 0, "message": "未找到科目2203的辅助余额数据"}

    # 构建 D3-2 行数据
    rows_data: list[dict] = []
    for aux_row in aux_rows[:_ROW_LIMIT]:
        rows_data.append({
            "rowId": str(uuid4()),
            "customerName": aux_row.aux_name or "",
            "companyCode": "",
            "nature": "其他",
            "relationType": "非关联方",
            "priorUnadjusted": float(aux_row.prior_balance),
            "priorAdjustment": 0,
            "priorReclass": 0,
            "agingPrior": {"within1": float(aux_row.prior_balance), "y1to2": 0, "y2to3": 0, "over3": 0},
            "debit": 0,
            "credit": 0,
            "entityReclass": 0,
            "endAje": 0,
            "endRje": 0,
            "agingAudited": {"within1": float(aux_row.current_balance), "y1to2": 0, "y2to3": 0, "over3": 0},
            "isConfirmed": "",
            "postPeriodSettlement": 0,
            "remark": "",
        })

    # 写入（merge模式：保留已有行、追加新客户）
    item_id = "D3-det-rows"
    existing_result = await db.execute(
        sa.text("SELECT remark FROM checklist_responses WHERE wp_id = :wp_id AND item_id = :item_id LIMIT 1"),
        {"wp_id": wp_id, "item_id": item_id},
    )
    existing_row = existing_result.fetchone()
    existing_rows: list[dict] = []
    if existing_row and existing_row.remark:
        try:
            existing_rows = json.loads(existing_row.remark)
        except (json.JSONDecodeError, TypeError):
            pass

    # Merge: 已有客户名不重复导入
    existing_names = {r.get("customerName", "") for r in existing_rows}
    new_rows = [r for r in rows_data if r["customerName"] not in existing_names]
    merged = existing_rows + new_rows

    remark_json = json.dumps(merged, ensure_ascii=False)
    await db.execute(
        sa.text("""
            INSERT INTO checklist_responses (id, wp_id, item_id, remark, updated_at)
            VALUES (:id, :wp_id, :item_id, :remark, NOW())
            ON CONFLICT (wp_id, item_id)
            DO UPDATE SET remark = :remark, updated_at = NOW()
        """),
        {"id": str(uuid4()), "wp_id": wp_id, "item_id": item_id, "remark": remark_json},
    )
    await db.commit()

    return {
        "ok": True,
        "imported_count": len(new_rows),
        "message": f"成功导入{len(new_rows)}行数据，{len(new_rows)}个新客户",
        "total_rows": len(merged),
    }


# ═══════════════════════════════════════════════════════════════════════════════
# 导出辅助函数
# ═══════════════════════════════════════════════════════════════════════════════


def _safe_float(val: Any) -> float:
    if val is None:
        return 0.0
    try:
        return float(val)
    except (ValueError, TypeError):
        return 0.0


def _safe_str(val: Any) -> str:
    if val is None:
        return ""
    return str(val).strip()


def _export_d3_2_row(data: dict) -> list:
    aging_prior = data.get("agingPrior", {})
    aging_audited = data.get("agingAudited", {})
    return [
        _safe_str(data.get("customerName")),
        _safe_str(data.get("companyCode")),
        _safe_str(data.get("nature")),
        _safe_str(data.get("relationType")),
        _safe_float(data.get("priorUnadjusted")),
        _safe_float(data.get("priorAdjustment")),
        _safe_float(data.get("priorReclass")),
        _safe_float(aging_prior.get("within1")),
        _safe_float(aging_prior.get("y1to2")),
        _safe_float(aging_prior.get("y2to3")),
        _safe_float(aging_prior.get("over3")),
        _safe_float(data.get("debit")),
        _safe_float(data.get("credit")),
        _safe_float(data.get("entityReclass")),
        _safe_float(data.get("endAje")),
        _safe_float(data.get("endRje")),
        _safe_float(aging_audited.get("within1")),
        _safe_float(aging_audited.get("y1to2")),
        _safe_float(aging_audited.get("y2to3")),
        _safe_float(aging_audited.get("over3")),
        _safe_str(data.get("isConfirmed")),
        _safe_float(data.get("postPeriodSettlement")),
        _safe_str(data.get("remark")),
    ]


def _export_d3_5_row(data: dict) -> list:
    # 键对齐前端 LongTermRow（reason/settlementAmount/disposalConclusion）；
    # 兼容历史键 unsettledReason/settledAmount。
    return [
        _safe_str(data.get("customerName")),
        _safe_float(data.get("endBalance")),
        _safe_str(data.get("aging")),
        _safe_str(data.get("businessDescription")),
        _safe_str(data.get("reason", data.get("unsettledReason"))),
        _safe_float(data.get("settlementAmount", data.get("settledAmount"))),
        _safe_str(data.get("disposalConclusion")),
        _safe_str(data.get("plan")),
        _safe_str(data.get("remark")),
    ]


def _export_d3_6_row(data: dict) -> list:
    return [
        _safe_str(data.get("partyName")),
        _safe_str(data.get("relationship")),
        _safe_float(data.get("priorBalance")),
        _safe_float(data.get("debit")),
        _safe_float(data.get("credit")),
        _safe_float(data.get("endBalance")),
        _safe_str(data.get("agingDescription")),
        _safe_str(data.get("nature")),
        _safe_str(data.get("indexRef")),
        _safe_str(data.get("remark")),
    ]


def _export_d3_7_row(data: dict) -> list:
    check_items = data.get("checkItems", [False] * 5)
    return [
        _safe_str(data.get("customerName")),
        _safe_str(data.get("date")),
        _safe_str(data.get("voucherNo")),
        _safe_str(data.get("businessContent")),
        _safe_str(data.get("counterAccount")),
        _safe_str(data.get("counterDetailAccount")),
        _safe_float(data.get("creditAmount")),
        _safe_str(data.get("supportingDoc")),
        *["Y" if c else "" for c in (check_items + [False] * 5)[:5]],
        _safe_str(data.get("indexRef")),
        _safe_str(data.get("isAbnormal")),
        _safe_str(data.get("remark")),
    ]


# ═══════════════════════════════════════════════════════════════════════════════
# 导入解析辅助函数
# ═══════════════════════════════════════════════════════════════════════════════


def _col_val(row: tuple, actual_headers: list[str], col_name: str) -> Any:
    try:
        idx = actual_headers.index(col_name)
        return row[idx] if idx < len(row) else None
    except (ValueError, IndexError):
        return None


def _parse_d3_2_row(
    row: tuple, actual_headers: list[str], segments: list[Any] | None = None
) -> dict:
    """解析 D3-2/F1-2 明细行。

    - segments 提供时（导入端点，Task 12.2）：账龄列按 label 动态匹配当前项目账龄配置，
      格式 `{label}(期初)` / `{label}(期末审定)`，写入 nested keyed agingPrior/agingAudited。
    - segments 为 None（legacy 纯函数/round-trip 测试）：读取旧固定列头 `期初审定账龄(1年以下)` 等，
      映射到固定 within1/y1to2/y2to3/over3 键。
    """
    base = {
        "rowId": str(uuid4()),
        "customerName": _safe_str(_col_val(row, actual_headers, "对方单位名称")),
        "companyCode": _safe_str(_col_val(row, actual_headers, "公司代码")),
        "nature": _safe_str(_col_val(row, actual_headers, "款项性质")) or "其他",
        "relationType": _safe_str(_col_val(row, actual_headers, "关联方类型")) or "非关联方",
        "priorUnadjusted": _safe_float(_col_val(row, actual_headers, "期初未审余额")),
        "priorAdjustment": _safe_float(_col_val(row, actual_headers, "期初账项调整")),
        "priorReclass": _safe_float(_col_val(row, actual_headers, "期初重分类调整")),
        "debit": _safe_float(_col_val(row, actual_headers, "借方发生")),
        "credit": _safe_float(_col_val(row, actual_headers, "贷方发生")),
        "entityReclass": _safe_float(_col_val(row, actual_headers, "被审计单位重分类调整")),
        "endAje": _safe_float(_col_val(row, actual_headers, "期末账项调整")),
        "endRje": _safe_float(_col_val(row, actual_headers, "期末重分类调整")),
        "isConfirmed": _safe_str(_col_val(row, actual_headers, "是否发函")),
        "postPeriodSettlement": _safe_float(_col_val(row, actual_headers, "期后结转")),
        "remark": _safe_str(_col_val(row, actual_headers, "备注")),
    }

    if segments:
        aging_nested, _unmatched = match_import_aging(
            lambda h: _col_val(row, actual_headers, h),
            actual_headers,
            segments,
            subject_aging_periods("D3"),
        )
        base.update(aging_nested)
    else:
        base["agingPrior"] = {
            "within1": _safe_float(_col_val(row, actual_headers, "期初审定账龄(1年以下)")),
            "y1to2": _safe_float(_col_val(row, actual_headers, "期初审定账龄(1~2年)")),
            "y2to3": _safe_float(_col_val(row, actual_headers, "期初审定账龄(2~3年)")),
            "over3": _safe_float(_col_val(row, actual_headers, "期初审定账龄(3年以上)")),
        }
        base["agingAudited"] = {
            "within1": _safe_float(_col_val(row, actual_headers, "审定账龄(1年以下)")),
            "y1to2": _safe_float(_col_val(row, actual_headers, "审定账龄(1~2年)")),
            "y2to3": _safe_float(_col_val(row, actual_headers, "审定账龄(2~3年)")),
            "over3": _safe_float(_col_val(row, actual_headers, "审定账龄(3年以上)")),
        }
    return base


def _parse_d3_5_row(row: tuple, actual_headers: list[str]) -> dict:
    # 键对齐前端 LongTermRow（reason/settlementAmount/disposalConclusion）
    return {
        "rowId": str(uuid4()),
        "customerName": _safe_str(_col_val(row, actual_headers, "对方单位名称")),
        "endBalance": _safe_float(_col_val(row, actual_headers, "期末余额")),
        "aging": _safe_str(_col_val(row, actual_headers, "账龄")),
        "businessDescription": _safe_str(_col_val(row, actual_headers, "经济业务说明")),
        "reason": _safe_str(_col_val(row, actual_headers, "未结转或未偿还的原因")),
        "settlementAmount": _safe_float(_col_val(row, actual_headers, "至审计日结转或偿还金额")),
        "disposalConclusion": _safe_str(_col_val(row, actual_headers, "处理结论")),
        "plan": _safe_str(_col_val(row, actual_headers, "处理计划")),
        "remark": _safe_str(_col_val(row, actual_headers, "备注")),
    }


def _parse_d3_6_row(row: tuple, actual_headers: list[str]) -> dict:
    return {
        "rowId": str(uuid4()),
        "partyName": _safe_str(_col_val(row, actual_headers, "关联方名称")),
        "relationship": _safe_str(_col_val(row, actual_headers, "关联关系")),
        "priorBalance": _safe_float(_col_val(row, actual_headers, "期初余额")),
        "debit": _safe_float(_col_val(row, actual_headers, "借方发生")),
        "credit": _safe_float(_col_val(row, actual_headers, "贷方发生")),
        "endBalance": _safe_float(_col_val(row, actual_headers, "期末余额")),
        "agingDescription": _safe_str(_col_val(row, actual_headers, "发生时间及账龄")),
        "nature": _safe_str(_col_val(row, actual_headers, "发生原因（款项性质）")),
        "indexRef": _safe_str(_col_val(row, actual_headers, "索引号")),
        "remark": _safe_str(_col_val(row, actual_headers, "备注")),
    }


def _parse_d3_7_row(row: tuple, actual_headers: list[str]) -> dict:
    check_items = [
        _safe_str(_col_val(row, actual_headers, f"核对内容{i}")).upper() in ("Y", "是", "TRUE", "1")
        for i in range(1, 6)
    ]
    return {
        "rowId": str(uuid4()),
        "customerName": _safe_str(_col_val(row, actual_headers, "客户名称")),
        "date": _safe_str(_col_val(row, actual_headers, "日期")),
        "voucherNo": _safe_str(_col_val(row, actual_headers, "凭证编号")),
        "businessContent": _safe_str(_col_val(row, actual_headers, "业务内容")),
        "counterAccount": _safe_str(_col_val(row, actual_headers, "对方科目")),
        "counterDetailAccount": _safe_str(_col_val(row, actual_headers, "对方明细科目")),
        "creditAmount": _safe_float(_col_val(row, actual_headers, "贷方金额")),
        "supportingDoc": _safe_str(_col_val(row, actual_headers, "支持性文件")),
        "checkItems": check_items,
        "indexRef": _safe_str(_col_val(row, actual_headers, "索引号")),
        "isAbnormal": _safe_str(_col_val(row, actual_headers, "是否异常")),
        "remark": _safe_str(_col_val(row, actual_headers, "备注说明")),
    }


def _parse_d3_3_row(row: tuple, actual_headers: list[str]) -> dict:
    return {
        "rowId": str(uuid4()),
        "description": _safe_str(_col_val(row, actual_headers, "调整事项说明")),
        "category": _safe_str(_col_val(row, actual_headers, "类别")) or "账项调整",
        "reportItem": _safe_str(_col_val(row, actual_headers, "报表项目")),
        "accountName": _safe_str(_col_val(row, actual_headers, "科目名称")),
        "noteItem": _safe_str(_col_val(row, actual_headers, "附注项目")),
        "placeholder": "",
        "debitAmount": _safe_float(_col_val(row, actual_headers, "借方调整金额")),
        "creditAmount": _safe_float(_col_val(row, actual_headers, "贷方调整金额")),
        "indexRef": _safe_str(_col_val(row, actual_headers, "索引")),
        "remark": _safe_str(_col_val(row, actual_headers, "备注")),
    }


def _parse_d3_4_row(row: tuple, actual_headers: list[str]) -> dict | None:
    row_key = _safe_str(_col_val(row, actual_headers, "行键"))
    label = _safe_str(_col_val(row, actual_headers, "项目"))
    if not row_key and not label:
        return None
    if row_key in ("debit-total", "debit-diff", "credit-total", "credit-diff"):
        return None
    return {
        "rowKey": row_key or label,
        "label": label or row_key,
        "amount": _safe_float(_col_val(row, actual_headers, "金额")),
        "source": _safe_str(_col_val(row, actual_headers, "来源")),
        "remark": _safe_str(_col_val(row, actual_headers, "备注")),
    }


def _export_d3_3_row(data: dict) -> list:
    return [
        _safe_str(data.get("description")),
        _safe_str(data.get("category")),
        _safe_str(data.get("reportItem")),
        _safe_str(data.get("accountName")),
        _safe_str(data.get("noteItem")),
        _safe_float(data.get("debitAmount")),
        _safe_float(data.get("creditAmount")),
        _safe_str(data.get("indexRef")),
        _safe_str(data.get("remark")),
    ]


def _export_d3_4_row(data: dict) -> list:
    return [
        _safe_str(data.get("rowKey")),
        _safe_str(data.get("label")),
        _safe_float(data.get("amount")),
        _safe_str(data.get("source")),
        _safe_str(data.get("remark")),
    ]


async def _fetch_response_map(
    wp_id: str,
    db: AsyncSession,
    prefix: str,
) -> dict[str, str]:
    import sqlalchemy as sa

    result = await db.execute(
        sa.text(
            "SELECT item_id, remark FROM checklist_responses "
            "WHERE wp_id = :wp_id AND item_id LIKE :prefix"
        ),
        {"wp_id": wp_id, "prefix": f"{prefix}%"},
    )
    return {row.item_id: (row.remark or "") for row in result.fetchall()}


async def _upsert_response(
    db: AsyncSession,
    wp_id: str,
    item_id: str,
    remark: str,
) -> None:
    import sqlalchemy as sa

    await db.execute(
        sa.text("""
            INSERT INTO checklist_responses (id, wp_id, item_id, remark, updated_at)
            VALUES (:id, :wp_id, :item_id, :remark, NOW())
            ON CONFLICT (wp_id, item_id)
            DO UPDATE SET remark = :remark, updated_at = NOW()
        """),
        {"id": str(uuid4()), "wp_id": wp_id, "item_id": item_id, "remark": remark},
    )


async def _stream_d3_1_export(
    wp_id: str,
    db: AsyncSession,
    headers: list[str],
    sheet: str,
) -> StreamingResponse:
    responses = await _fetch_response_map(wp_id, db, "D3-adj-")

    wb = Workbook()
    ws = wb.active
    ws.title = sheet
    ws.append(headers)
    ws.freeze_panes = "A2"

    for section, row_key, label in _D3_1_ROWS:
        row_values = [
            _D3_1_SECTION_LABELS[section],
            row_key,
            label,
        ]
        for field_key, _header, is_text in _D3_1_FIELDS:
            item_id = f"D3-adj-{section}-{row_key}-{field_key}"
            raw = responses.get(item_id, "")
            if is_text:
                row_values.append(raw)
            else:
                row_values.append(_safe_float(raw) if raw else 0.0)
        ws.append(row_values)

    buffer = io.BytesIO()
    wb.save(buffer)
    buffer.seek(0)

    from urllib.parse import quote
    encoded_filename = quote(f"{sheet}_数据.xlsx")
    return StreamingResponse(
        buffer,
        media_type="application/vnd.openxmlformats-officedocument.spreadsheetml.sheet",
        headers={"Content-Disposition": f"attachment; filename*=UTF-8''{encoded_filename}"},
    )


def _parse_d3_1_row(row: tuple, actual_headers: list[str]) -> dict | None:
    section_label = _safe_str(_col_val(row, actual_headers, "区块"))
    row_key = _safe_str(_col_val(row, actual_headers, "行键"))
    label = _safe_str(_col_val(row, actual_headers, "项目"))
    if not row_key and not label:
        return None

    section = _D3_1_SECTION_BY_LABEL.get(section_label, "")
    if not section and row_key:
        for sec, rk, _lbl in _D3_1_ROWS:
            if rk == row_key:
                section = sec
                break
    if not row_key:
        for sec, rk, lbl in _D3_1_ROWS:
            if lbl == label:
                section = sec
                row_key = rk
                break
    if not section or not row_key:
        return None

    parsed: dict[str, Any] = {"section": section, "rowKey": row_key, "label": label}
    for field_key, header, is_text in _D3_1_FIELDS:
        val = _col_val(row, actual_headers, header)
        parsed[field_key] = _safe_str(val) if is_text else _safe_float(val)
    return parsed


async def _persist_d3_1_rows(
    wp_id: str,
    rows_data: list[dict],
    db: AsyncSession,
) -> int:
    field_count = 0
    for row in rows_data:
        section = row.get("section", "")
        row_key = row.get("rowKey", "")
        if not section or not row_key:
            continue
        for field_key, _header, is_text in _D3_1_FIELDS:
            val = row.get(field_key)
            if val is None or (not is_text and val == ""):
                continue
            item_id = f"D3-adj-{section}-{row_key}-{field_key}"
            remark = str(val) if is_text else str(_safe_float(val))
            await _upsert_response(db, wp_id, item_id, remark)
            field_count += 1
    return field_count

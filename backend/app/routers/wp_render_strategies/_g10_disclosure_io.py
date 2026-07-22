"""G10 附注披露 v2 store ↔ 多工作表 Excel（上市/国企）."""

from __future__ import annotations

import io
import re
from datetime import datetime
from typing import Any

from openpyxl import Workbook, load_workbook
from openpyxl.styles import Font

from ._cycle_import_export_common import safe_float, safe_str

_LISTED_MV_ROWS: list[tuple[str, str]] = [
    ("mv_trading_bond", "    其中：发行的交易性债券"),
    ("mv_derivative", "衍生金融负债"),
    ("mv_other", "其他"),
    ("mv_designated_bond", "    其中：债券"),
    ("mv_designated_other", "其他"),
]

_SOE_BAL_ROWS: list[tuple[str, str]] = [
    ("soe_trading_bond", "    其中：发行的交易性债券"),
    ("soe_derivative", "衍生金融负债"),
    ("soe_other", "其他"),
    ("soe_designated_other", "其他"),
]

_MV_HEADERS = ["行键", "项目", "期初余额", "本期增加", "本期减少", "期末余额"]
_MV_KEYS = ["rowKey", "label", "openingAmount", "increaseAmount", "decreaseAmount", "closingAmount"]

_BAL_HEADERS = ["行键", "项目", "期末公允价值", "期初公允价值"]
_BAL_KEYS = ["rowKey", "label", "currentAmount", "priorAmount"]

_DESIGNATED_HEADERS = ["行键", "项目", "期初余额", "期末余额", "指定的理由和依据"]
_DESIGNATED_KEYS = ["rowKey", "label", "openingAmount", "closingAmount", "designationReason"]

_FV_KEYS = ["rowKey", "label", "fvChangeAmount", "creditRiskCurrent", "creditRiskCumulative"]

_DERIV_HEADERS = ["行键", "项目", "期末余额", "上年年末余额"]
_DERIV_KEYS = ["rowKey", "label", "currentAmount", "priorAmount"]

_NOTE_HEADERS = ["字段", "内容"]
_NOTE_KEYS = ["field", "content"]


def _resolve_audit_year(store: dict[str, Any]) -> int | None:
    y = store.get("auditYear")
    if y is None:
        return None
    try:
        n = int(y)
        return n if 1900 <= n <= 2100 else None
    except (TypeError, ValueError):
        return None


def _fv_headers(audit_year: int | None) -> list[str]:
    y = f"{audit_year}年" if audit_year else "本年"
    return [
        "行键",
        "项目",
        f"{y}公允价值变动额",
        "因自身信用风险变动引起的公允价值本年变动额",
        "因自身信用风险变动引起的公允价值累计变动额",
    ]


def _empty_movement() -> dict[str, float]:
    return {
        "openingAmount": 0.0,
        "increaseAmount": 0.0,
        "decreaseAmount": 0.0,
        "closingAmount": 0.0,
    }


def _empty_balance() -> dict[str, float]:
    return {"currentAmount": 0.0, "priorAmount": 0.0}


def default_listed_store() -> dict[str, Any]:
    return {
        "version": 2,
        "movement": {k: _empty_movement() for k, _ in _LISTED_MV_ROWS},
        "designatedDetail": {
            "designated_1": {
                "label": "发行的普通债权",
                "openingAmount": 0.0,
                "closingAmount": 0.0,
                "designationReason": "",
            },
        },
        "fvCreditRisk": {
            "fv_1": {
                "label": "发行的普通债权",
                "fvChangeAmount": 0.0,
                "creditRiskCurrent": 0.0,
                "creditRiskCumulative": 0.0,
            },
        },
        "derivativeRows": [],
        "derivativeNote": "",
        "maturityDiffNote": "",
    }


def default_soe_store() -> dict[str, Any]:
    return {
        "version": 2,
        "balance": {k: _empty_balance() for k, _ in _SOE_BAL_ROWS},
        "fvCreditRisk": {
            "fv_1": {
                "label": "发行的普通债权",
                "fvChangeAmount": 0.0,
                "creditRiskCurrent": 0.0,
                "creditRiskCumulative": 0.0,
            },
        },
        "maturityDiffNote": "",
    }


def _normalize_listed(payload: Any) -> dict[str, Any]:
    base = default_listed_store()
    if not isinstance(payload, dict) or payload.get("version") != 2:
        return base
    audit_year = _resolve_audit_year(payload)
    movement = dict(base["movement"])
    for k, v in (payload.get("movement") or {}).items():
        if isinstance(v, dict):
            movement[k] = {**_empty_movement(), **v}
    designated = {}
    for k, v in (payload.get("designatedDetail") or {}).items():
        if isinstance(v, dict):
            designated[k] = {
                "label": safe_str(v.get("label")),
                "openingAmount": safe_float(v.get("openingAmount")),
                "closingAmount": safe_float(v.get("closingAmount")),
                "designationReason": safe_str(v.get("designationReason")),
            }
    fv = {}
    for k, v in (payload.get("fvCreditRisk") or {}).items():
        if isinstance(v, dict):
            fv[k] = {
                "label": safe_str(v.get("label")),
                "fvChangeAmount": safe_float(v.get("fvChangeAmount")),
                "creditRiskCurrent": safe_float(v.get("creditRiskCurrent")),
                "creditRiskCumulative": safe_float(v.get("creditRiskCumulative")),
            }
    deriv = []
    for row in payload.get("derivativeRows") or []:
        if isinstance(row, dict) and safe_str(row.get("rowKey")):
            deriv.append({
                "rowKey": safe_str(row.get("rowKey")),
                "label": safe_str(row.get("label")),
                "currentAmount": safe_float(row.get("currentAmount")),
                "priorAmount": safe_float(row.get("priorAmount")),
            })
    return {
        "version": 2,
        "auditYear": audit_year,
        "movement": movement or base["movement"],
        "designatedDetail": designated or base["designatedDetail"],
        "fvCreditRisk": fv or base["fvCreditRisk"],
        "derivativeRows": deriv,
        "derivativeNote": safe_str(payload.get("derivativeNote")),
        "maturityDiffNote": safe_str(payload.get("maturityDiffNote")),
    }


def _normalize_soe(payload: Any) -> dict[str, Any]:
    base = default_soe_store()
    if not isinstance(payload, dict) or payload.get("version") != 2:
        return base
    audit_year = _resolve_audit_year(payload)
    balance = dict(base["balance"])
    for k, v in (payload.get("balance") or {}).items():
        if isinstance(v, dict):
            balance[k] = {**_empty_balance(), **v}
    fv = {}
    for k, v in (payload.get("fvCreditRisk") or {}).items():
        if isinstance(v, dict):
            fv[k] = {
                "label": safe_str(v.get("label")),
                "fvChangeAmount": safe_float(v.get("fvChangeAmount")),
                "creditRiskCurrent": safe_float(v.get("creditRiskCurrent")),
                "creditRiskCumulative": safe_float(v.get("creditRiskCumulative")),
            }
    return {
        "version": 2,
        "auditYear": audit_year,
        "balance": balance or base["balance"],
        "fvCreditRisk": fv or base["fvCreditRisk"],
        "maturityDiffNote": safe_str(payload.get("maturityDiffNote")),
    }


def _append_sheet(
    wb: Workbook,
    title: str,
    headers: list[str],
    rows: list[list[Any]],
) -> None:
    ws = wb.create_sheet(title)
    ws.append(headers)
    for r in rows:
        ws.append(r)
    ws.freeze_panes = "A2"
    ws["A1"].font = Font(bold=True)


def _row_dict(headers: list[str], row: tuple) -> dict[str, Any]:
    out: dict[str, Any] = {}
    for i, h in enumerate(headers):
        out[h] = row[i] if i < len(row) else None
    return out


def _find_header_row(ws) -> tuple[int, list[str]]:
    for i, row in enumerate(ws.iter_rows(min_row=1, max_row=10, values_only=True), start=1):
        cells = [safe_str(c) for c in row if c is not None]
        if any(c in ("行键", "rowKey", "字段", "field") for c in cells):
            headers = [safe_str(c) for c in row]
            return i, headers
    return 1, [safe_str(c) for c in next(ws.iter_rows(min_row=1, max_row=1, values_only=True))]


def _parse_table(ws, expected_headers: list[str], field_keys: list[str]) -> list[dict[str, Any]]:
    if ws is None:
        return []
    header_row, headers = _find_header_row(ws)
    rows: list[dict[str, Any]] = []
    for row in ws.iter_rows(min_row=header_row + 1, values_only=True):
        if not any(c is not None and safe_str(c) for c in row):
            continue
        d = _row_dict(headers, row)
        label = safe_str(d.get("项目") or d.get("label"))
        if label in ("合  计", "合计"):
            continue
        key = safe_str(d.get("行键") or d.get("rowKey"))
        if not key and not label:
            continue
        parsed: dict[str, Any] = {"rowKey": key or f"row_{len(rows) + 1}"}
        for hdr, fk in zip(expected_headers, field_keys):
            val = d.get(hdr)
            if fk in ("rowKey", "label", "designationReason", "field", "content"):
                parsed[fk] = safe_str(val)
            else:
                parsed[fk] = safe_float(val)
        if not parsed.get("label") and label:
            parsed["label"] = label
        rows.append(parsed)
    return rows


def _map_fv_column_indices(headers: list[str]) -> dict[str, int]:
    idx: dict[str, int] = {}
    for i, h in enumerate(headers):
        s = safe_str(h)
        if not s:
            continue
        if "行键" in s or s.lower() == "rowkey":
            idx["rowKey"] = i
        elif "项目" in s:
            idx["label"] = i
        elif "公允价值变动" in s and "信用" not in s:
            idx["fvChangeAmount"] = i
        elif ("本年" in s or "信用风险本年" in s) and "信用" in s:
            idx["creditRiskCurrent"] = i
        elif ("累计" in s or "信用风险累计" in s) and "信用" in s:
            idx["creditRiskCumulative"] = i
    return idx


def _parse_fv_table(ws) -> list[dict[str, Any]]:
    if ws is None:
        return []
    header_row, headers = _find_header_row(ws)
    col = _map_fv_column_indices(headers)
    if "rowKey" not in col and "label" not in col:
        return []
    rows: list[dict[str, Any]] = []
    for row in ws.iter_rows(min_row=header_row + 1, values_only=True):
        if not any(c is not None and safe_str(c) for c in row):
            continue
        values = list(row)
        label = safe_str(values[col["label"]]) if "label" in col and col["label"] < len(values) else ""
        if label in ("合  计", "合计"):
            continue
        key = safe_str(values[col["rowKey"]]) if "rowKey" in col and col["rowKey"] < len(values) else ""
        if not key and not label:
            continue
        parsed: dict[str, Any] = {
            "rowKey": key or f"fv_{len(rows) + 1}",
            "label": label,
            "fvChangeAmount": safe_float(values[col["fvChangeAmount"]]) if "fvChangeAmount" in col and col["fvChangeAmount"] < len(values) else 0.0,
            "creditRiskCurrent": safe_float(values[col["creditRiskCurrent"]]) if "creditRiskCurrent" in col and col["creditRiskCurrent"] < len(values) else 0.0,
            "creditRiskCumulative": safe_float(values[col["creditRiskCumulative"]]) if "creditRiskCumulative" in col and col["creditRiskCumulative"] < len(values) else 0.0,
        }
        rows.append(parsed)
    return rows


def _infer_audit_year_from_fv_header(headers: list[str]) -> int | None:
    for h in headers:
        m = re.search(r"(20\d{2})年", safe_str(h))
        if m:
            return int(m.group(1))
    return None


def build_listed_disclosure_workbook(payload: Any, *, template_only: bool = False) -> Workbook:
    store = default_listed_store() if template_only else _normalize_listed(payload)
    audit_year = _resolve_audit_year(store) or (datetime.now().year if template_only else None)
    fv_headers = _fv_headers(audit_year)
    wb = Workbook()
    wb.remove(wb.active)

    mv_rows = []
    for key, default_label in _LISTED_MV_ROWS:
        cell = store["movement"].get(key, _empty_movement())
        label = default_label
        mv_rows.append([
            key, label,
            cell.get("openingAmount", 0),
            cell.get("increaseAmount", 0),
            cell.get("decreaseAmount", 0),
            cell.get("closingAmount", 0),
        ])
    _append_sheet(wb, "变动表", _MV_HEADERS, mv_rows)

    des_rows = []
    for key, cell in store["designatedDetail"].items():
        des_rows.append([
            key,
            cell.get("label") or "发行的普通债权",
            cell.get("openingAmount", 0),
            cell.get("closingAmount", 0),
            cell.get("designationReason", ""),
        ])
    _append_sheet(wb, "指定明细", _DESIGNATED_HEADERS, des_rows)

    fv_rows = []
    for key, cell in store["fvCreditRisk"].items():
        fv_rows.append([
            key,
            cell.get("label") or "发行的普通债权",
            cell.get("fvChangeAmount", 0),
            cell.get("creditRiskCurrent", 0),
            cell.get("creditRiskCumulative", 0),
        ])
    _append_sheet(wb, "信用风险", fv_headers, fv_rows)

    deriv_rows = []
    for row in store["derivativeRows"]:
        deriv_rows.append([
            row.get("rowKey", ""),
            row.get("label", ""),
            row.get("currentAmount", 0),
            row.get("priorAmount", 0),
        ])
    _append_sheet(wb, "衍生负债", _DERIV_HEADERS, deriv_rows)

    note_rows: list[list[Any]] = []
    if audit_year:
        note_rows.append(["auditYear", audit_year])
    note_rows.extend([
        ["derivativeNote", store.get("derivativeNote", "")],
        ["maturityDiffNote", store.get("maturityDiffNote", "")],
    ])
    _append_sheet(wb, "说明", _NOTE_HEADERS, note_rows)
    return wb


def build_soe_disclosure_workbook(payload: Any, *, template_only: bool = False) -> Workbook:
    store = default_soe_store() if template_only else _normalize_soe(payload)
    audit_year = _resolve_audit_year(store) or (datetime.now().year if template_only else None)
    fv_headers = _fv_headers(audit_year)
    wb = Workbook()
    wb.remove(wb.active)

    bal_rows = []
    for key, default_label in _SOE_BAL_ROWS:
        cell = store["balance"].get(key, _empty_balance())
        bal_rows.append([
            key, default_label,
            cell.get("currentAmount", 0),
            cell.get("priorAmount", 0),
        ])
    _append_sheet(wb, "余额表", _BAL_HEADERS, bal_rows)

    fv_rows = []
    for key, cell in store["fvCreditRisk"].items():
        fv_rows.append([
            key,
            cell.get("label") or "发行的普通债权",
            cell.get("fvChangeAmount", 0),
            cell.get("creditRiskCurrent", 0),
            cell.get("creditRiskCumulative", 0),
        ])
    _append_sheet(wb, "信用风险", fv_headers, fv_rows)

    note_rows: list[list[Any]] = []
    if audit_year:
        note_rows.append(["auditYear", audit_year])
    note_rows.append(["maturityDiffNote", store.get("maturityDiffNote", "")])
    _append_sheet(wb, "说明", _NOTE_HEADERS, note_rows)
    return wb


def parse_listed_disclosure_workbook(content: bytes) -> tuple[dict[str, Any], list[str], int]:
    errors: list[str] = []
    try:
        wb = load_workbook(io.BytesIO(content), read_only=True, data_only=True)
    except Exception as e:
        return default_listed_store(), [f"无法解析 xlsx: {e}"], 0

    store = default_listed_store()
    count = 0

    mv = _parse_table(wb["变动表"] if "变动表" in wb.sheetnames else None, _MV_HEADERS, _MV_KEYS)
    if mv:
        movement: dict[str, Any] = {}
        for row in mv:
            key = row["rowKey"]
            movement[key] = {
                "openingAmount": row["openingAmount"],
                "increaseAmount": row["increaseAmount"],
                "decreaseAmount": row["decreaseAmount"],
                "closingAmount": row["closingAmount"],
            }
        store["movement"] = movement
        count += len(mv)

    des = _parse_table(
        wb["指定明细"] if "指定明细" in wb.sheetnames else None,
        _DESIGNATED_HEADERS, _DESIGNATED_KEYS,
    )
    if des:
        designated: dict[str, Any] = {}
        for row in des:
            key = row["rowKey"]
            designated[key] = {
                "label": row.get("label", ""),
                "openingAmount": row["openingAmount"],
                "closingAmount": row["closingAmount"],
                "designationReason": row.get("designationReason", ""),
            }
        store["designatedDetail"] = designated
        count += len(des)

    fv_ws = wb["信用风险"] if "信用风险" in wb.sheetnames else None
    fv = _parse_fv_table(fv_ws)
    if fv:
        fv_store: dict[str, Any] = {}
        for row in fv:
            key = row["rowKey"]
            fv_store[key] = {
                "label": row.get("label", ""),
                "fvChangeAmount": row["fvChangeAmount"],
                "creditRiskCurrent": row["creditRiskCurrent"],
                "creditRiskCumulative": row["creditRiskCumulative"],
            }
        store["fvCreditRisk"] = fv_store
        count += len(fv)
        if fv_ws is not None:
            _, hdrs = _find_header_row(fv_ws)
            inferred = _infer_audit_year_from_fv_header(hdrs)
            if inferred:
                store["auditYear"] = inferred

    deriv = _parse_table(
        wb["衍生负债"] if "衍生负债" in wb.sheetnames else None,
        _DERIV_HEADERS, _DERIV_KEYS,
    )
    if deriv:
        store["derivativeRows"] = [
            {
                "rowKey": row["rowKey"],
                "label": row.get("label", ""),
                "currentAmount": row["currentAmount"],
                "priorAmount": row["priorAmount"],
            }
            for row in deriv
        ]
        count += len(deriv)

    notes = _parse_table(wb["说明"] if "说明" in wb.sheetnames else None, _NOTE_HEADERS, _NOTE_KEYS)
    for row in notes:
        field = row.get("field", "")
        content_val = row.get("content", "")
        if field == "auditYear":
            y = _resolve_audit_year({"auditYear": content_val})
            if y:
                store["auditYear"] = y
        elif field in ("derivativeNote", "maturityDiffNote"):
            store[field] = content_val

    if count == 0:
        errors.append("未识别到有效数据行，请使用平台模板（含「变动表」等工作表）")
    return store, errors, count


def parse_soe_disclosure_workbook(content: bytes) -> tuple[dict[str, Any], list[str], int]:
    errors: list[str] = []
    try:
        wb = load_workbook(io.BytesIO(content), read_only=True, data_only=True)
    except Exception as e:
        return default_soe_store(), [f"无法解析 xlsx: {e}"], 0

    store = default_soe_store()
    count = 0

    bal = _parse_table(wb["余额表"] if "余额表" in wb.sheetnames else None, _BAL_HEADERS, _BAL_KEYS)
    if bal:
        balance: dict[str, Any] = {}
        for row in bal:
            key = row["rowKey"]
            balance[key] = {
                "currentAmount": row["currentAmount"],
                "priorAmount": row["priorAmount"],
            }
        store["balance"] = balance
        count += len(bal)

    fv_ws = wb["信用风险"] if "信用风险" in wb.sheetnames else None
    fv = _parse_fv_table(fv_ws)
    if fv:
        fv_store: dict[str, Any] = {}
        for row in fv:
            key = row["rowKey"]
            fv_store[key] = {
                "label": row.get("label", ""),
                "fvChangeAmount": row["fvChangeAmount"],
                "creditRiskCurrent": row["creditRiskCurrent"],
                "creditRiskCumulative": row["creditRiskCumulative"],
            }
        store["fvCreditRisk"] = fv_store
        count += len(fv)
        if fv_ws is not None:
            _, hdrs = _find_header_row(fv_ws)
            inferred = _infer_audit_year_from_fv_header(hdrs)
            if inferred:
                store["auditYear"] = inferred

    notes = _parse_table(wb["说明"] if "说明" in wb.sheetnames else None, _NOTE_HEADERS, _NOTE_KEYS)
    for row in notes:
        field = row.get("field", "")
        content_val = row.get("content", "")
        if field == "auditYear":
            y = _resolve_audit_year({"auditYear": content_val})
            if y:
                store["auditYear"] = y
        elif field == "maturityDiffNote":
            store["maturityDiffNote"] = content_val

    if count == 0:
        errors.append("未识别到有效数据行，请使用平台模板（含「余额表」等工作表）")
    return store, errors, count

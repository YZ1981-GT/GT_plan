"""G7 长期股权投资(main组) — G7-2/G7-3 导入导出.

G7-2 以原底稿的三条业务链为数据模型：
成本法(A:AI)、权益法(A:BB)、减值准备(A:U)，并保留未审、AJE、RJE、审定四层。
导入同时支持系统三业务区工作簿和原始 ``明细表G7-2`` 模板。
"""

from __future__ import annotations

import io
from typing import Any
from uuid import uuid4

from fastapi import APIRouter, Body, Depends, File, HTTPException, Query, UploadFile
from fastapi.responses import StreamingResponse
from openpyxl import Workbook, load_workbook
from openpyxl.styles import Font
from openpyxl.utils import get_column_letter
from sqlalchemy.ext.asyncio import AsyncSession

from app.core.database import get_db
from app.deps import get_current_user
from app.models.core import User

from ._cycle_import_export_common import (
    ROW_LIMIT,
    build_workbook_template,
    export_row_by_keys,
    load_json_rows,
    parse_row_by_headers,
    parse_upload_xlsx,
    safe_float,
    safe_str,
    upsert_json_rows,
    workbook_to_response,
)

router = APIRouter(tags=["g7-main-import-export"])

# ═══════════════════════════════════════════════════════════════════════════════
# G7-2 明细表 — 成本法 / 权益法 / 减值准备
# ═══════════════════════════════════════════════════════════════════════════════

_COMMON_FIELDS = [
    ("记录ID", "id"),
    ("序号", "seq"),
    ("被投资单位名称", "investeeName"),
    ("初始投资成本", "initialInvestmentCost"),
    ("投资比例", "investmentRatio"),
    ("投资日期", "investmentDate"),
    ("投资方式", "investmentMethod"),
]

_COST_FIELDS = _COMMON_FIELDS + [
    ("本期现金股利", "cashDividend"),
    ("未审期初比例", "openingRatio"),
    ("未审期初金额", "openingAmount"),
    ("未审增加比例", "increaseRatio"),
    ("未审增加金额", "increaseAmount"),
    ("增加索引", "increaseIndex"),
    ("未审减少比例", "decreaseRatio"),
    ("未审减少金额", "decreaseAmount"),
    ("减少索引", "decreaseIndex"),
    ("未审期末比例", "closingRatio"),
    ("未审期末金额", "closingAmount"),
    ("期初AJE", "openingAje"),
    ("期初RJE", "openingRje"),
    ("本期增加AJE", "ajeIncrease"),
    ("本期减少AJE", "ajeDecrease"),
    ("本期增加RJE", "rjeIncrease"),
    ("本期减少RJE", "rjeDecrease"),
    ("审定期初比例", "auditedOpeningRatio"),
    ("审定期初金额", "auditedOpeningAmount"),
    ("审定增加比例", "auditedIncreaseRatio"),
    ("审定增加金额", "auditedIncreaseAmount"),
    ("审定减少比例", "auditedDecreaseRatio"),
    ("审定减少金额", "auditedDecreaseAmount"),
    ("审定期末比例", "auditedClosingRatio"),
    ("审定期末金额", "auditedClosingAmount"),
]

_EQUITY_FIELDS = _COMMON_FIELDS + [
    ("投资关系", "relationship"),
    ("未审期初比例", "openingRatio"),
    ("未审期初金额", "openingAmount"),
    ("未审增加比例", "increaseRatio"),
    ("投资成本增加", "costIncrease"),
    ("损益调整", "profitLossAdjustment"),
    ("其他综合收益", "otherComprehensiveIncome"),
    ("其他权益变动", "otherEquityChange"),
    ("权益增加小计", "equityIncreaseSubtotal"),
    ("其他增加", "otherIncrease"),
    ("未审减少比例", "decreaseRatio"),
    ("投资成本减少", "costDecrease"),
    ("本期分回利润", "dividendReceived"),
    ("其他减少", "otherDecrease"),
    ("未审期末比例", "closingRatio"),
    ("未审期末金额", "closingAmount"),
    ("期初AJE", "openingAje"),
    ("期初RJE", "openingRje"),
    ("AJE-成本增加", "ajeCostIncrease"),
    ("AJE-损益调整", "ajeProfitLoss"),
    ("AJE-其他综合收益", "ajeOci"),
    ("AJE-其他权益变动", "ajeOtherEquity"),
    ("AJE-其他增加", "ajeOtherIncrease"),
    ("AJE-成本减少", "ajeCostDecrease"),
    ("AJE-分回利润", "ajeDividend"),
    ("AJE-其他减少", "ajeOtherDecrease"),
    ("RJE-成本增加", "rjeCostIncrease"),
    ("RJE-损益调整", "rjeProfitLoss"),
    ("RJE-其他综合收益", "rjeOci"),
    ("RJE-其他权益变动", "rjeOtherEquity"),
    ("RJE-其他增加", "rjeOtherIncrease"),
    ("RJE-成本减少", "rjeCostDecrease"),
    ("RJE-分回利润", "rjeDividend"),
    ("RJE-其他减少", "rjeOtherDecrease"),
    ("审定期初比例", "auditedOpeningRatio"),
    ("审定期初金额", "auditedOpeningAmount"),
    ("审定增加比例", "auditedIncreaseRatio"),
    ("审定成本增加", "auditedCostIncrease"),
    ("审定损益调整", "auditedProfitLoss"),
    ("审定其他综合收益", "auditedOci"),
    ("审定其他权益变动", "auditedOtherEquity"),
    ("审定权益增加小计", "auditedEquityIncreaseSubtotal"),
    ("审定其他增加", "auditedOtherIncrease"),
    ("审定减少比例", "auditedDecreaseRatio"),
    ("审定成本减少", "auditedCostDecrease"),
    ("审定分回利润", "auditedDividend"),
    ("审定其他减少", "auditedOtherDecrease"),
    ("审定期末比例", "auditedClosingRatio"),
    ("审定期末金额", "auditedClosingAmount"),
]

_IMPAIRMENT_FIELDS = _COMMON_FIELDS + [
    ("来源记录ID", "sourceId"),
    ("投资关系", "relationship"),
    ("未审期初", "openingAmount"),
    ("未审增加", "increaseAmount"),
    ("未审减少", "decreaseAmount"),
    ("未审期末", "closingAmount"),
    ("备注/索引", "remark"),
    ("期初AJE", "openingAje"),
    ("期初RJE", "openingRje"),
    ("增加AJE", "ajeIncrease"),
    ("减少AJE", "ajeDecrease"),
    ("增加RJE", "rjeIncrease"),
    ("减少RJE", "rjeDecrease"),
    ("审定期初", "auditedOpeningAmount"),
    ("审定增加", "auditedIncreaseAmount"),
    ("审定减少", "auditedDecreaseAmount"),
    ("审定期末", "auditedClosingAmount"),
]

_G7_2_SECTIONS = [
    ("成本法", "cost", _COST_FIELDS),
    ("权益法", "equity", _EQUITY_FIELDS),
    ("减值准备", "impairment", _IMPAIRMENT_FIELDS),
]

_RATIO_KEYS = {
    "investmentRatio", "openingRatio", "increaseRatio", "decreaseRatio", "closingRatio",
    "auditedOpeningRatio", "auditedIncreaseRatio", "auditedDecreaseRatio", "auditedClosingRatio",
}
_TEXT_KEYS = {
    "id", "section", "investeeName", "investmentDate", "investmentMethod", "increaseIndex",
    "decreaseIndex", "relationship", "sourceId", "remark",
}

# ═══════════════════════════════════════════════════════════════════════════════
# G7-3 调整分录汇总（导出含扩展列；导入 CORE 兼容旧模板）
# ═══════════════════════════════════════════════════════════════════════════════

_G7_3_CORE_HEADERS = [
    "调整事项说明", "类别（报表调整/账项调整/其他）", "报表项目", "科目名称",
    "附注项目", "……", "借方调整金额", "贷方调整金额", "索引", "备注",
]
_G7_3_CORE_KEYS = [
    "description", "category", "reportItem", "accountName", "noteItem",
    "_spacer", "debitAmount", "creditAmount", "indexRef", "remark",
]
_G7_3_OPTIONAL_HEADERS = ("科目代码", "来源标记", "被投资单位")
_G7_3_HEADERS = [
    "调整事项说明", "类别（报表调整/账项调整/其他）", "报表项目", "科目名称", "科目代码",
    "附注项目", "……", "借方调整金额", "贷方调整金额", "索引", "备注", "来源标记", "被投资单位",
]
_G7_3_KEYS = [
    "description", "category", "reportItem", "accountName", "accountCode",
    "noteItem", "_spacer", "debitAmount", "creditAmount", "indexRef", "remark",
    "sourceKind", "investeeName",
]

# ═══════════════════════════════════════════════════════════════════════════════
# 通用辅助
# ═══════════════════════════════════════════════════════════════════════════════

_ITEM_IDS = {
    "G7-2": "G7-2-rows",
    "G7-3": "G7-3-rows",
}

_SUPPORTED_SHEETS = set(_ITEM_IDS.keys())


def _validate_sheet(sheet: str) -> None:
    if sheet not in _SUPPORTED_SHEETS:
        raise HTTPException(400, f"不支持的sheet: {sheet}。支持: {sorted(_SUPPORTED_SHEETS)}")


# ═══════════════════════════════════════════════════════════════════════════════
# G7-2 三业务区导出与原模板导入
# ═══════════════════════════════════════════════════════════════════════════════

def _build_g7_2_multi_sheet_workbook(rows: list[dict], *, template_only: bool = False) -> Workbook:
    """按成本法、权益法、减值准备三个原表业务区导出。"""
    wb = Workbook()
    wb.remove(wb.active)  # type: ignore[arg-type]

    for sheet_name, section, fields in _G7_2_SECTIONS:
        headers = [item[0] for item in fields]
        keys = [item[1] for item in fields]
        ws = wb.create_sheet(title=sheet_name)
        ws.append([f"G7-2 长期股权投资明细表 — {sheet_name}"])
        ws.merge_cells(start_row=1, start_column=1, end_row=1, end_column=len(headers))
        ws["A1"].font = Font(bold=True, size=12)
        ws.append(headers)
        ws.freeze_panes = "A3"
        for col_idx, key in enumerate(keys, start=1):
            ws.column_dimensions[get_column_letter(col_idx)].width = (
                20 if key in {"investeeName", "remark"} else 15
            )

        if not template_only:
            for row_data in (row for row in rows if row.get("section") == section):
                ws.append(export_row_by_keys(row_data, keys))
                excel_row = ws.max_row
                for col_idx, key in enumerate(keys, start=1):
                    if key in _RATIO_KEYS:
                        ws.cell(excel_row, col_idx).number_format = "0.00%"

    ws_guide = wb.create_sheet("编制说明")
    ws_guide.append(["G7-2 长期股权投资明细表 编制说明"])
    ws_guide.append([])
    guidance_lines = [
        "本模板按原始底稿三大区块拆分：成本法、权益法、减值准备。",
        "每个区块保留未审数、账项调整(AJE)、重分类调整(RJE)、审定数四层。",
        "比例使用Excel百分比格式，例如直接输入51%（底层数值为0.51）。",
        "蓝色/审定结果字段由系统导入后重算；请勿用审定结果替代未审及调整来源。",
        "成本法审定期末 = 审定期初 + 审定增加 - 审定减少。",
        "权益增加小计 = 损益调整 + 其他综合收益 + 其他权益变动。",
        "权益法审定期末按原表口径，不包含“其他增加/其他减少”栏。",
        "减值审定期末 = 审定期初 + 审定增加 - 审定减少。",
        "也可直接导入项目原始《G7 长期股权投资.xlsx》中的“明细表G7-2”。",
    ]
    for line in guidance_lines:
        ws_guide.append([line])
    ws_guide.column_dimensions["A"].width = 80
    return wb


def _ratio(value: Any) -> float:
    if isinstance(value, str) and value.strip().endswith("%"):
        return safe_float(value.strip()[:-1]) / 100
    return safe_float(value)


def _text(value: Any) -> str:
    if value is None:
        return ""
    if hasattr(value, "isoformat"):
        return value.isoformat()
    return safe_str(value)


def _typed_value(key: str, value: Any) -> Any:
    if key == "relationship":
        relation = _text(value)
        return {
            "子公司": "subsidiary",
            "合营企业": "joint_venture",
            "合营": "joint_venture",
            "联营企业": "associate",
            "联营": "associate",
        }.get(relation, relation)
    if key in _TEXT_KEYS:
        return _text(value)
    if key in _RATIO_KEYS:
        return _ratio(value)
    return safe_float(value)


def _parse_structured_g7_2(wb: Any) -> tuple[list[dict], list[str]]:
    rows: list[dict] = []
    errors: list[str] = []
    for sheet_name, section, fields in _G7_2_SECTIONS:
        if sheet_name not in wb.sheetnames:
            errors.append(f"缺少工作表：{sheet_name}")
            continue
        ws = wb[sheet_name]
        headers = [item[0] for item in fields]
        key_by_header = dict(fields)
        actual = [_text(cell.value) for cell in ws[2]]
        missing = [name for name in ("序号", "被投资单位名称") if name not in actual]
        if missing:
            errors.append(f"工作表[{sheet_name}]缺少列：{', '.join(missing)}")
            continue
        for values in ws.iter_rows(min_row=3, values_only=True):
            value_by_header = {
                header: values[index] if index < len(values) else None
                for index, header in enumerate(actual)
            }
            investee = _text(value_by_header.get("被投资单位名称"))
            if not investee or investee == "……":
                continue
            parsed: dict[str, Any] = {"section": section}
            for header in headers:
                key = key_by_header[header]
                parsed[key] = _typed_value(key, value_by_header.get(header))
            parsed["id"] = parsed.get("id") or str(uuid4())
            parsed["section"] = section
            rows.append(parsed)
            if len(rows) >= ROW_LIMIT:
                errors.append(f"数据行超过{ROW_LIMIT}行限制，已截断")
                return rows, errors
    return rows, errors


def _base_original_row(
    ws: Any,
    row_no: int,
    *,
    section: str,
    name_col: int,
    initial_col: int,
    ratio_col: int,
    date_col: int,
    method_col: int,
) -> dict[str, Any] | None:
    name = _text(ws.cell(row_no, name_col).value)
    if not name or name == "……":
        return None
    return {
        "id": str(uuid4()),
        "section": section,
        "seq": safe_float(ws.cell(row_no, 1).value),
        "investeeName": name,
        "initialInvestmentCost": safe_float(ws.cell(row_no, initial_col).value),
        "investmentRatio": _ratio(ws.cell(row_no, ratio_col).value),
        "investmentDate": _text(ws.cell(row_no, date_col).value),
        "investmentMethod": _text(ws.cell(row_no, method_col).value),
    }


def _parse_original_g7_2(wb: Any) -> tuple[list[dict], list[str]]:
    """按原始模板固定坐标读取 10条成本法、5条合营、5条联营及对应减值。"""
    ws = wb["明细表G7-2"]
    rows: list[dict] = []
    sources: dict[tuple[str, int], dict] = {}

    for index, row_no in enumerate(range(16, 26), start=1):
        row = _base_original_row(
            ws, row_no, section="cost", name_col=2, initial_col=6,
            ratio_col=8, date_col=9, method_col=10,
        )
        if row is None:
            continue
        row.update({
            "cashDividend": safe_float(ws.cell(row_no, 11).value),
            "openingRatio": _ratio(ws.cell(row_no, 12).value),
            "openingAmount": safe_float(ws.cell(row_no, 13).value),
            "increaseRatio": _ratio(ws.cell(row_no, 14).value),
            "increaseAmount": safe_float(ws.cell(row_no, 15).value),
            "increaseIndex": _text(ws.cell(row_no, 16).value),
            "decreaseRatio": _ratio(ws.cell(row_no, 17).value),
            "decreaseAmount": safe_float(ws.cell(row_no, 18).value),
            "decreaseIndex": _text(ws.cell(row_no, 19).value),
            "closingRatio": _ratio(ws.cell(row_no, 20).value),
            "closingAmount": safe_float(ws.cell(row_no, 21).value),
            "openingAje": safe_float(ws.cell(row_no, 22).value),
            "openingRje": safe_float(ws.cell(row_no, 23).value),
            "ajeIncrease": safe_float(ws.cell(row_no, 24).value),
            "ajeDecrease": safe_float(ws.cell(row_no, 25).value),
            "rjeIncrease": safe_float(ws.cell(row_no, 26).value),
            "rjeDecrease": safe_float(ws.cell(row_no, 27).value),
        })
        rows.append(row)
        sources[("subsidiary", index)] = row

    equity_ranges = [
        ("joint_venture", range(33, 38)),
        ("associate", range(40, 45)),
    ]
    for relationship, row_range in equity_ranges:
        for index, row_no in enumerate(row_range, start=1):
            row = _base_original_row(
                ws, row_no, section="equity", name_col=2, initial_col=3,
                ratio_col=4, date_col=5, method_col=6,
            )
            if row is None:
                continue
            row.update({
                "relationship": relationship,
                "openingRatio": _ratio(ws.cell(row_no, 7).value),
                "openingAmount": safe_float(ws.cell(row_no, 8).value),
                "increaseRatio": _ratio(ws.cell(row_no, 9).value),
                "costIncrease": safe_float(ws.cell(row_no, 10).value),
                "profitLossAdjustment": safe_float(ws.cell(row_no, 11).value),
                "otherComprehensiveIncome": safe_float(ws.cell(row_no, 12).value),
                "otherEquityChange": safe_float(ws.cell(row_no, 13).value),
                "equityIncreaseSubtotal": safe_float(ws.cell(row_no, 14).value),
                "otherIncrease": safe_float(ws.cell(row_no, 15).value),
                "decreaseRatio": _ratio(ws.cell(row_no, 16).value),
                "costDecrease": safe_float(ws.cell(row_no, 17).value),
                "dividendReceived": safe_float(ws.cell(row_no, 18).value),
                "otherDecrease": safe_float(ws.cell(row_no, 19).value),
                "closingRatio": _ratio(ws.cell(row_no, 20).value),
                "closingAmount": safe_float(ws.cell(row_no, 21).value),
                "openingAje": safe_float(ws.cell(row_no, 22).value),
                "openingRje": safe_float(ws.cell(row_no, 23).value),
                "ajeCostIncrease": safe_float(ws.cell(row_no, 24).value),
                "ajeProfitLoss": safe_float(ws.cell(row_no, 25).value),
                "ajeOci": safe_float(ws.cell(row_no, 26).value),
                "ajeOtherEquity": safe_float(ws.cell(row_no, 27).value),
                "ajeOtherIncrease": safe_float(ws.cell(row_no, 28).value),
                "ajeCostDecrease": safe_float(ws.cell(row_no, 29).value),
                "ajeDividend": safe_float(ws.cell(row_no, 30).value),
                "ajeOtherDecrease": safe_float(ws.cell(row_no, 31).value),
                "rjeCostIncrease": safe_float(ws.cell(row_no, 32).value),
                "rjeProfitLoss": safe_float(ws.cell(row_no, 33).value),
                "rjeOci": safe_float(ws.cell(row_no, 34).value),
                "rjeOtherEquity": safe_float(ws.cell(row_no, 35).value),
                "rjeOtherIncrease": safe_float(ws.cell(row_no, 36).value),
                "rjeCostDecrease": safe_float(ws.cell(row_no, 37).value),
                "rjeDividend": safe_float(ws.cell(row_no, 38).value),
                "rjeOtherDecrease": safe_float(ws.cell(row_no, 39).value),
            })
            rows.append(row)
            sources[(relationship, index)] = row

    impairment_ranges = [
        ("subsidiary", range(53, 63)),
        ("joint_venture", range(66, 71)),
        ("associate", range(73, 78)),
    ]
    for relationship, row_range in impairment_ranges:
        for index, row_no in enumerate(row_range, start=1):
            source = sources.get((relationship, index))
            if source is None:
                continue
            rows.append({
                "id": str(uuid4()),
                "section": "impairment",
                "sourceId": source["id"],
                "seq": index,
                "investeeName": source["investeeName"],
                "relationship": relationship,
                "initialInvestmentCost": source["initialInvestmentCost"],
                "investmentRatio": source["investmentRatio"],
                "investmentDate": source["investmentDate"],
                "investmentMethod": source["investmentMethod"],
                "openingAmount": safe_float(ws.cell(row_no, 7).value),
                "increaseAmount": safe_float(ws.cell(row_no, 8).value),
                "decreaseAmount": safe_float(ws.cell(row_no, 9).value),
                "closingAmount": safe_float(ws.cell(row_no, 10).value),
                "remark": _text(ws.cell(row_no, 11).value),
                "openingAje": safe_float(ws.cell(row_no, 12).value),
                "openingRje": safe_float(ws.cell(row_no, 13).value),
                "ajeIncrease": safe_float(ws.cell(row_no, 14).value),
                "ajeDecrease": safe_float(ws.cell(row_no, 15).value),
                "rjeIncrease": safe_float(ws.cell(row_no, 16).value),
                "rjeDecrease": safe_float(ws.cell(row_no, 17).value),
            })
    return rows, []


def _parse_g7_2_import(content: bytes) -> tuple[list[dict], list[str]]:
    wb = load_workbook(io.BytesIO(content), read_only=True, data_only=True)
    try:
        if "明细表G7-2" in wb.sheetnames:
            return _parse_original_g7_2(wb)
        if any(name in wb.sheetnames for name, _, _ in _G7_2_SECTIONS):
            return _parse_structured_g7_2(wb)
        return [], ["无法识别G7-2结构：请使用原始“明细表G7-2”或系统导出的三业务区模板"]
    finally:
        wb.close()


def _validate_g7_2_rows(rows: list[dict]) -> list[dict[str, Any]]:
    """复核前端派生金额，防止绕过前端后保存失真的审定链。"""
    errors: list[dict[str, Any]] = []

    def check(row: dict, field: str, expected: float) -> None:
        actual = safe_float(row.get(field))
        variance = round(actual - expected, 2)
        if abs(variance) > 0.01:
            errors.append({
                "rowId": safe_str(row.get("id")),
                "investeeName": safe_str(row.get("investeeName")),
                "field": field,
                "expected": round(expected, 2),
                "actual": round(actual, 2),
                "variance": variance,
            })

    for row in rows[:ROW_LIMIT]:
        section = row.get("section")
        opening = safe_float(row.get("openingAmount"))
        increase = safe_float(row.get("increaseAmount"))
        decrease = safe_float(row.get("decreaseAmount"))
        opening_aje = safe_float(row.get("openingAje"))
        opening_rje = safe_float(row.get("openingRje"))

        if section == "cost":
            check(row, "closingAmount", opening + increase - decrease)
            audited_opening = opening + opening_aje + opening_rje
            audited_increase = (
                increase + safe_float(row.get("ajeIncrease")) + safe_float(row.get("rjeIncrease"))
            )
            audited_decrease = (
                decrease + safe_float(row.get("ajeDecrease")) + safe_float(row.get("rjeDecrease"))
            )
            check(row, "auditedOpeningAmount", audited_opening)
            check(row, "auditedIncreaseAmount", audited_increase)
            check(row, "auditedDecreaseAmount", audited_decrease)
            check(row, "auditedClosingAmount", audited_opening + audited_increase - audited_decrease)
        elif section == "equity":
            equity_subtotal = (
                safe_float(row.get("profitLossAdjustment"))
                + safe_float(row.get("otherComprehensiveIncome"))
                + safe_float(row.get("otherEquityChange"))
            )
            check(row, "equityIncreaseSubtotal", equity_subtotal)
            unadjusted_closing = (
                opening
                + safe_float(row.get("costIncrease"))
                + equity_subtotal
                + safe_float(row.get("otherIncrease"))
                - safe_float(row.get("costDecrease"))
                - safe_float(row.get("dividendReceived"))
                - safe_float(row.get("otherDecrease"))
            )
            check(row, "closingAmount", unadjusted_closing)
            audited_opening = opening + opening_aje + opening_rje
            audited_cost_increase = (
                safe_float(row.get("costIncrease"))
                + safe_float(row.get("ajeCostIncrease"))
                + safe_float(row.get("rjeCostIncrease"))
            )
            audited_equity_increase = sum(
                safe_float(row.get(key))
                for key in (
                    "profitLossAdjustment", "ajeProfitLoss", "rjeProfitLoss",
                    "otherComprehensiveIncome", "ajeOci", "rjeOci",
                    "otherEquityChange", "ajeOtherEquity", "rjeOtherEquity",
                )
            )
            audited_cost_decrease = (
                safe_float(row.get("costDecrease"))
                + safe_float(row.get("ajeCostDecrease"))
                + safe_float(row.get("rjeCostDecrease"))
            )
            audited_dividend = (
                safe_float(row.get("dividendReceived"))
                + safe_float(row.get("ajeDividend"))
                + safe_float(row.get("rjeDividend"))
            )
            check(
                row,
                "auditedClosingAmount",
                audited_opening + audited_cost_increase + audited_equity_increase
                - audited_cost_decrease - audited_dividend,
            )
        elif section == "impairment":
            check(row, "closingAmount", opening + increase - decrease)
            audited_opening = opening + opening_aje + opening_rje
            audited_increase = (
                increase + safe_float(row.get("ajeIncrease")) + safe_float(row.get("rjeIncrease"))
            )
            audited_decrease = (
                decrease + safe_float(row.get("ajeDecrease")) + safe_float(row.get("rjeDecrease"))
            )
            check(row, "auditedClosingAmount", audited_opening + audited_increase - audited_decrease)
    return errors


# ═══════════════════════════════════════════════════════════════════════════════
# 端点
# ═══════════════════════════════════════════════════════════════════════════════

@router.post("/api/workpapers/{wp_id}/g7-main/validate-detail")
async def g7_main_validate_detail(
    wp_id: str,
    rows: list[dict] = Body(...),
    current_user: User = Depends(get_current_user),
) -> dict[str, Any]:
    """校验G7-2三业务区公式链；wp_id用于路由归属和权限上下文。"""
    del wp_id, current_user
    errors = _validate_g7_2_rows(rows)
    return {"ok": not errors, "errors": errors, "checked_count": min(len(rows), ROW_LIMIT)}


@router.post("/api/workpapers/{wp_id}/g7-main/export-template")
async def g7_main_export_template(
    wp_id: str,
    sheet: str = Query(...),
    current_user: User = Depends(get_current_user),
) -> StreamingResponse:
    """导出空模板：G7-2(三业务区多sheet) / G7-3(单sheet)。"""
    _validate_sheet(sheet)

    if sheet == "G7-2":
        wb = _build_g7_2_multi_sheet_workbook([], template_only=True)
        return workbook_to_response(wb, "G7-2_明细表_模板.xlsx")
    else:  # G7-3
        wb = build_workbook_template(
            "G7-3",
            _G7_3_HEADERS,
            title="G7-3 调整分录汇总",
            guidance=[
                "G7-3 调整分录 编制说明",
                "",
                "类别填写：账项调整、报表调整或其他；同一调整事项的借贷金额必须相等。",
                "科目名称请使用标准科目名称；可另填科目代码。来源标记/被投资单位用于建议草稿与按单位回写。",
                "确认后同步调整分录模块，并按 1511/1512 回写 G7-1。",
            ],
        )
        return workbook_to_response(wb, "G7-3_调整分录_模板.xlsx")


@router.post("/api/workpapers/{wp_id}/g7-main/export-data")
async def g7_main_export_data(
    wp_id: str,
    sheet: str = Query(...),
    db: AsyncSession = Depends(get_db),
    current_user: User = Depends(get_current_user),
) -> StreamingResponse:
    """导出数据：G7-2(三业务区多sheet含数据) / G7-3(单sheet含数据)。"""
    _validate_sheet(sheet)
    item_id = _ITEM_IDS[sheet]
    rows = await load_json_rows(db, wp_id, item_id, field="conclusion")

    if sheet == "G7-2":
        wb = _build_g7_2_multi_sheet_workbook(rows, template_only=False)
        return workbook_to_response(wb, "G7-2_明细表_数据.xlsx")
    else:  # G7-3
        wb = build_workbook_template(
            "G7-3",
            _G7_3_HEADERS,
            title="G7-3 调整分录汇总",
            guidance=["类别填写账项调整、报表调整或其他；同一调整事项借贷金额必须相等。"],
        )
        ws = wb["G7-3"]
        for d in rows:
            ws.append(export_row_by_keys(d, _G7_3_KEYS))
        return workbook_to_response(wb, "G7-3_调整分录_数据.xlsx")


@router.post("/api/workpapers/{wp_id}/g7-main/import-data")
async def g7_main_import_data(
    wp_id: str,
    sheet: str = Query(...),
    file: UploadFile = File(...),
    db: AsyncSession = Depends(get_db),
    current_user: User = Depends(get_current_user),
) -> dict[str, Any]:
    """导入数据：G7-2(系统三业务区或原始工作簿) / G7-3(单sheet)。"""
    _validate_sheet(sheet)
    if not file.filename or not file.filename.endswith(".xlsx"):
        raise HTTPException(400, "请上传 .xlsx 格式文件")
    content = await file.read()
    if len(content) > 10 * 1024 * 1024:
        raise HTTPException(400, "文件大小不能超过10MB")

    item_id = _ITEM_IDS[sheet]
    errors: list[str] = []
    rows: list[dict] = []
    warnings: list[str] = []

    if sheet == "G7-2":
        rows, errors = _parse_g7_2_import(content)
    else:  # G7-3
        try:
            actual, raw = parse_upload_xlsx(
                content,
                _G7_3_CORE_HEADERS,
                header_row=2,
                require_all_headers=True,
            )
        except ValueError as e:
            return {"ok": False, "errors": [str(e)], "imported_count": 0}
        except Exception:
            raise HTTPException(400, "无法解析xlsx文件")
        missing_opt = [h for h in _G7_3_OPTIONAL_HEADERS if h not in actual]
        if missing_opt:
            warnings.append(f"兼容旧模板：缺列 {', '.join(missing_opt)} 已按空值导入")
        for i, r in enumerate(raw, start=1):
            if i > ROW_LIMIT:
                errors.append(f"数据行超过{ROW_LIMIT}行限制，已截断")
                break
            rows.append(parse_row_by_headers(
                r, actual, _G7_3_KEYS, expected_headers=_G7_3_HEADERS,
            ))

    if errors and not rows:
        return {"ok": False, "errors": errors, "imported_count": 0}

    await upsert_json_rows(db, wp_id, item_id, rows, field="conclusion")
    out: dict[str, Any] = {"ok": True, "imported_count": len(rows), "errors": errors}
    if warnings:
        out["warning"] = "；".join(warnings)
    if len(rows) >= ROW_LIMIT:
        trunc = f"数据行数超过{ROW_LIMIT}行限制，已截断"
        out["warning"] = f"{out.get('warning')}；{trunc}" if out.get("warning") else trunc
    return out

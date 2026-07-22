"""G11 投资收益 — 导入导出（G11-1~5 + 附注上市/国企）."""

from __future__ import annotations

import io
import json
from typing import Any
from uuid import uuid4

import sqlalchemy as sa

from fastapi import APIRouter, Depends, File, HTTPException, Query, UploadFile
from fastapi.responses import StreamingResponse
from openpyxl import Workbook, load_workbook
from openpyxl.styles import Font
from openpyxl.utils import get_column_letter
from sqlalchemy.ext.asyncio import AsyncSession

from app.core.database import get_db
from app.deps import get_current_user
from app.models.core import User

from ._cycle_import_export_common import (
    ROW_LIMIT,
    build_workbook_template,
    export_row_by_keys,
    import_rows_generic,
    is_numeric_field_key,
    load_json_payload,
    load_json_rows,
    parse_row_by_headers,
    parse_upload_xlsx,
    safe_float,
    safe_str,
    upsert_json_payload,
    upsert_json_rows,
    workbook_to_response,
)
from ._g11_disclosure_io import (
    build_listed_disclosure_workbook,
    build_soe_disclosure_workbook,
    parse_listed_disclosure_workbook,
    parse_soe_disclosure_workbook,
)

_G11_1_HEADERS = [
    "行键", "项目", "分组",
    "本期未审数", "本期调整", "上期未审数", "上期调整",
    "原因分析", "索引号",
]
_G11_1_KEYS = [
    "rowKey", "label", "group",
    "currentUnadjusted", "currentAdjustment", "priorUnadjusted", "priorAdjustment",
    "reasonAnalysis", "indexRef",
]

_G11_2_HEADERS = [
    "序号", "行键", "项目", "投资类型", "被投资单位", "处置子类",
    "本期未审数", "本期调整", "本期审定数", "本期占比",
    "上年未审数", "上期调整", "上年审定数", "上期占比",
    "变动额", "变动原因/索引号",
]
_G11_2_KEYS = [
    "seq", "rowKey", "itemName", "group", "investeeName", "tradingDisposeSubtype",
    "currentUnadjusted", "currentAdjustment", "currentAudited", "currentShare",
    "priorUnadjusted", "priorAdjustment", "priorAudited", "priorShare",
    "changeAmount", "reasonIndex",
]

_G11_3_HEADERS = [
    "调整事项说明", "类别", "报表项目", "科目代码", "科目名称", "附注项目",
    "回写行", "借方调整金额", "贷方调整金额", "索引", "备注",
]
_G11_3_KEYS = [
    "description", "category", "reportItem", "accountCode", "accountName", "noteItem",
    "adjudicationRowKey", "debitAmount", "creditAmount", "indexRef", "remark",
]

_G11_4_HEADERS = [
    "行键", "项目名称", "本期发生额", "本期期初余额", "本期期末余额",
    "上期审定数", "上期期初余额", "上期期末余额", "市场收益率", "异常说明",
]
_G11_4_KEYS = [
    "rowKey", "itemName", "currentIncome", "currentOpening", "currentClosing",
    "priorAudited", "priorOpening", "priorClosing", "marketYield", "abnormalNote",
]

_G11_5_SEGMENTS: list[tuple[str, list[str], list[str]]] = [
    ("凭证基础", ["序号", "日期", "凭证编号", "业务内容", "对方科目", "对方明细科目", "贷方金额", "附件"],
     ["seq", "voucherDate", "voucherNo", "businessContent", "counterAccount", "counterDetail", "creditAmount", "attachment"]),
    ("核对内容", ["序号", "支持性文件描述", "核对1完整", "核对2授权", "核对3金额正确", "核对4期间恰当", "核对5账务正确"],
     ["seq", "supportingDocDesc", "check1", "check2", "check3", "check4", "check5"]),
    ("结论", ["序号", "索引号", "是否异常", "异常说明", "风险等级", "备注", "来源"],
     ["seq", "indexNo", "isAbnormal", "abnormalDesc", "riskLevel", "remark", "source"]),
]

_G11_5_FLAT_HEADERS = [
    "序号", "日期", "凭证编号", "业务内容", "对方科目", "对方明细科目", "贷方金额", "附件",
    "支持性文件描述", "核对1完整", "核对2授权", "核对3金额正确", "核对4期间恰当", "核对5账务正确",
    "索引号", "是否异常", "异常说明", "风险等级", "备注", "来源",
]
_G11_5_FLAT_KEYS = [
    "seq", "voucherDate", "voucherNo", "businessContent", "counterAccount", "counterDetail", "creditAmount", "attachment",
    "supportingDocDesc", "check1", "check2", "check3", "check4", "check5",
    "indexNo", "isAbnormal", "abnormalDesc", "riskLevel", "remark", "source",
]

_G11_SPECS: dict[str, dict[str, Any]] = {
    "G11-1": {
        "item_id": "G11-adj-rows",
        "title": "G11-1 投资收益审定表",
        "headers": _G11_1_HEADERS,
        "field_keys": _G11_1_KEYS,
        "guidance": [
            "G11-1 审定表 编制说明",
            "",
            "审定数=未审+调整；导入后由前端重算变动额/变动率。",
        ],
        "storage_shape": "object",
    },
    "G11-2": {
        "item_id": "G11-detail-rows",
        "title": "G11-2 投资收益明细分析表",
        "headers": _G11_2_HEADERS,
        "field_keys": _G11_2_KEYS,
        "allow_missing_headers": True,
        "header_aliases": {
            "变动原因/索引号": ["变动原因索引号", "变动原因/索引", "reasonIndex"],
            "行键": ["rowKey"],
            "投资类型": ["分组", "group"],
            "处置子类": ["tradingDisposeSubtype", "子类"],
        },
        "guidance": [
            "G11-2 明细分析表 编制说明",
            "",
            "与 G11-1 同口径分项；审定数=未审+调整；占比/变动率导入后由前端重算。",
            "|变动率|>20% 须填变动原因/索引。",
            "处置子类（股票/债券/非套期衍生/套期衍生/其他）用于上市附注「处置交易性」明细灌数。",
        ],
    },
    "G11-3": {
        "item_id": "G11-aje-rows",
        "title": "G11-3 调整分录汇总",
        "headers": _G11_3_HEADERS,
        "field_keys": _G11_3_KEYS,
        "allow_missing_headers": True,
        "header_aliases": {
            "调整事项说明": ["摘要", "调整说明", "事项说明"],
            "类别": ["分录类型", "类型", "entryType"],
            "借方调整金额": ["借方金额", "借方"],
            "贷方调整金额": ["贷方金额", "贷方"],
            "索引": ["索引号", "indexRef"],
            "回写行": ["审定表行", "adjudicationRowKey"],
            "科目代码": ["科目编码"],
        },
        "guidance": [
            "G11-3 调整分录汇总 编制说明",
            "",
            "列对齐 Excel：调整事项说明/类别/报表项目/科目/附注/借贷/索引/备注。",
            "类别：账项调整（回写审定）/报表调整（仅列报）/其他。",
            "仅 6111 贷−借净额按「回写行」分项写入 G11-1/G11-2；整表借贷须平衡。",
            "兼容旧模板列名：分录类型→类别、摘要→调整事项说明。",
        ],
    },
    "G11-4": {
        "item_id": "G11-return-rate-rows",
        "title": "G11-4 收益率分析表",
        "headers": _G11_4_HEADERS,
        "field_keys": _G11_4_KEYS,
        "guidance": [
            "G11-4 收益率分析",
            "",
            "平均投资②=(期初+期末)/2；比率③=发生额①/平均投资；变动⑦=③−⑥；|变动|>5pp须填异常说明。",
            "行键与 G11-1 对齐时可从前端「从 G11-1 带入」；市场收益率列可选填外部对标。",
        ],
    },
    "G11-5": {
        "item_id": "G11-voucher-rows",
        "title": "G11-5 凭证检查表",
        "headers": _G11_5_FLAT_HEADERS,
        "field_keys": _G11_5_FLAT_KEYS,
        "guidance": [
            "G11-5 凭证检查",
            "",
            "导出含3区段工作表；导入支持宽表或3区段。核对列填 ✓ / ✗ / 未测（三态）。",
            "样本选取参数存 G11-vc-params（测试总体/特定样本/抽样方法/本期发生额等）；",
            "检查比例=已查贷方合计÷本期发生额；可从G11-1带入；金额/账务异常可推送 G11-3。",
            "截止样本回填默认未测；跨期强制异常。",
        ],
    },
    "附注上市": {
        "item_id": "G11-disclosure-listed",
        "title": "附注披露信息（上市公司）",
        "storage_field": "remark",
        "build_workbook": lambda payload, template_only=False: build_listed_disclosure_workbook(
            payload, template_only=template_only,
        ),
        "parse_import": parse_listed_disclosure_workbook,
        "guidance": [
            "附注披露（上市公司）编制说明",
            "",
            "多工作表：主表（项目/本期/上期）+ 处置交易性明细 + 说明。",
            "建议先在 G11-1/G11-2 维护分项，再在附注页「分项带入」。",
            "主表合计应与 G11-1 审定（6111）勾稽。",
        ],
    },
    "附注国企": {
        "item_id": "G11-disclosure-soe",
        "title": "附注披露信息（国企）",
        "storage_field": "remark",
        "build_workbook": lambda payload, template_only=False: build_soe_disclosure_workbook(
            payload, template_only=template_only,
        ),
        "parse_import": parse_soe_disclosure_workbook,
        "guidance": [
            "附注披露（国企）编制说明",
            "",
            "多工作表：主表（项目/本期/上期）+ 说明（汇回重大限制）。",
            "建议「分项带入」后按实际裁剪不存在项目。",
        ],
    },
}

_SUPPORTED = set(_G11_SPECS.keys())
router = APIRouter(tags=["g11-import-export"])
_HEADER_ROW = 2


def _parse_g11_1_row(row: tuple, headers: list[str]) -> dict[str, Any]:
    values = list(row) + [None] * max(0, len(headers) - len(row))
    out: dict[str, Any] = {}
    for h, key in zip(_G11_1_HEADERS, _G11_1_KEYS):
        if h not in headers:
            continue
        idx = headers.index(h)
        raw = values[idx] if idx < len(values) else None
        if is_numeric_field_key(key):
            out[key] = safe_float(raw)
        else:
            out[key] = safe_str(raw)
    return out


def _rows_to_g11_1_store(rows: list[dict]) -> dict[str, dict]:
    store: dict[str, dict] = {}
    for row in rows:
        key = safe_str(row.get("rowKey"))
        if not key:
            continue
        store[key] = {
            "currentUnadjusted": safe_float(row.get("currentUnadjusted")),
            "currentAdjustment": safe_float(row.get("currentAdjustment")),
            "priorUnadjusted": safe_float(row.get("priorUnadjusted")),
            "priorAdjustment": safe_float(row.get("priorAdjustment")),
            "reasonAnalysis": safe_str(row.get("reasonAnalysis")),
            "indexRef": safe_str(row.get("indexRef")),
            "label": safe_str(row.get("label")),
            "group": safe_str(row.get("group")),
        }
    return store


async def _load_g11_adj_remark(db: AsyncSession, wp_id: str) -> dict[str, dict]:
    result = await db.execute(
        sa.text(
            "SELECT remark FROM checklist_responses WHERE wp_id = :wp_id AND item_id = :iid LIMIT 1"
        ),
        {"wp_id": wp_id, "iid": "G11-adj-rows"},
    )
    row = result.fetchone()
    if not row or not row.remark:
        return {}
    try:
        parsed = json.loads(row.remark)
        if isinstance(parsed, dict):
            return parsed
        if isinstance(parsed, list):
            return _rows_to_g11_1_store(parsed)
    except (json.JSONDecodeError, TypeError):
        pass
    return {}


async def _upsert_g11_adj_store(db: AsyncSession, wp_id: str, store: dict[str, dict]) -> None:
    proj = await db.execute(
        sa.text("SELECT project_id FROM working_paper WHERE id = :wp_id AND is_deleted = false"),
        {"wp_id": wp_id},
    )
    project_id = proj.scalar_one_or_none()
    if not project_id:
        raise ValueError(f"working_paper not found: {wp_id}")
    await db.execute(
        sa.text("""
            INSERT INTO checklist_responses (id, project_id, wp_id, item_id, remark, updated_at, created_at)
            VALUES (:id, :project_id, :wp_id, :item_id, :payload, NOW(), NOW())
            ON CONFLICT (wp_id, item_id)
            DO UPDATE SET remark = :payload, updated_at = NOW()
        """),
        {
            "id": str(uuid4()),
            "project_id": str(project_id),
            "wp_id": wp_id,
            "item_id": "G11-adj-rows",
            "payload": json.dumps(store, ensure_ascii=False),
        },
    )
    await db.commit()


def _export_g11_1_rows(store: dict[str, dict]) -> list[list[Any]]:
    out: list[list[Any]] = []
    for key, fields in store.items():
        if not isinstance(fields, dict):
            continue
        row = {
            "rowKey": key,
            "label": fields.get("label") or key,
            "group": fields.get("group") or "",
            **fields,
        }
        out.append(export_row_by_keys(row, _G11_1_KEYS))
    return out


def _validate(sheet: str) -> None:
    if sheet not in _SUPPORTED:
        raise HTTPException(400, f"不支持的sheet: {sheet}。支持: {sorted(_SUPPORTED)}")


def _parse_check_val(raw: Any) -> Any:
    """三态：✓/通过→True；✗/不通过→False；未测/空→None。"""
    if raw is None:
        return None
    if isinstance(raw, bool):
        return raw
    s = safe_str(raw).strip().lower()
    if not s or s in {"未测", "n/a", "na", "-", "—", "null"}:
        return None
    if s in {"✓", "√", "是", "y", "yes", "true", "1", "通过"}:
        return True
    if s in {"✗", "×", "否", "n", "no", "false", "0", "不通过"}:
        return False
    return None


def _format_check_export(val: Any) -> str:
    if val is True or val in ("✓", "是", "true", "1", True):
        return "✓"
    if val is False or val in ("✗", "否", "false", "0", False):
        return "✗"
    return "未测"


def _parse_g11_5_row(row: tuple, headers: list[str]) -> dict[str, Any]:
    values = list(row) + [None] * max(0, len(headers) - len(row))
    out: dict[str, Any] = {"id": str(uuid4())}
    check_keys = {"check1", "check2", "check3", "check4", "check5"}
    for h, key in zip(_G11_5_FLAT_HEADERS, _G11_5_FLAT_KEYS):
        if h not in headers:
            continue
        idx = headers.index(h)
        raw = values[idx] if idx < len(values) else None
        if key in check_keys:
            out[key] = _parse_check_val(raw)
        elif key == "isAbnormal":
            out[key] = safe_str(raw) in {"是", "✓", "true", "1", "Y", "y"}
        elif is_numeric_field_key(key):
            out[key] = safe_float(raw)
        else:
            out[key] = safe_str(raw)
    return out


def _build_g11_5_workbook(rows: list[dict], *, template_only: bool = False) -> Workbook:
    wb = Workbook()
    wb.remove(wb.active)
    for seg_name, seg_headers, seg_keys in _G11_5_SEGMENTS:
        ws = wb.create_sheet(title=seg_name)
        ws.append([f"G11-5 凭证检查 — {seg_name}"])
        ws.merge_cells(start_row=1, start_column=1, end_row=1, end_column=max(len(seg_headers), 1))
        ws["A1"].font = Font(bold=True, size=12)
        ws.append(seg_headers)
        ws.freeze_panes = "A3"
        for col_idx in range(1, len(seg_headers) + 1):
            ws.column_dimensions[get_column_letter(col_idx)].width = 14
        if not template_only:
            for i, row_data in enumerate(rows, start=1):
                merged = {**row_data, "seq": row_data.get("seq", i)}
                row_out = export_row_by_keys(merged, seg_keys)
                if seg_name == "核对内容":
                    for j, k in enumerate(seg_keys):
                        if k.startswith("check"):
                            row_out[j] = _format_check_export(row_out[j])
                if seg_name == "结论" and "isAbnormal" in seg_keys:
                    idx = seg_keys.index("isAbnormal")
                    val = merged.get("isAbnormal")
                    row_out[idx] = "是" if val in (True, "是", "✓") else "否"
                ws.append(row_out)
    ws_guide = wb.create_sheet("编制说明")
    for line in _G11_SPECS["G11-5"]["guidance"]:
        ws_guide.append([line])
    ws_guide.column_dimensions["A"].width = 80
    return wb


def _parse_g11_5_upload(content: bytes) -> tuple[list[dict], list[str]]:
    wb = load_workbook(io.BytesIO(content), read_only=True, data_only=True)
    errors: list[str] = []
    rows_dict: dict[int, dict] = {}

    if any(seg in n for n in wb.sheetnames for seg in ("凭证基础", "核对内容", "结论")):
        for seg_name, seg_headers, seg_keys in _G11_5_SEGMENTS:
            ws = next((wb[n] for n in wb.sheetnames if seg_name in n), None)
            if ws is None:
                errors.append(f"缺少工作表: {seg_name}")
                continue
            for row_idx, row in enumerate(ws.iter_rows(min_row=3, values_only=True)):
                if all(v is None for v in row):
                    continue
                if row_idx >= ROW_LIMIT:
                    errors.append(f"数据行超过{ROW_LIMIT}行限制，已截断")
                    break
                if row_idx not in rows_dict:
                    rows_dict[row_idx] = {"id": str(uuid4())}
                values = list(row) + [None] * max(0, len(seg_headers) - len(row))
                for col_i, key in enumerate(seg_keys):
                    raw = values[col_i] if col_i < len(values) else None
                    if key in {"check1", "check2", "check3", "check4", "check5"}:
                        rows_dict[row_idx][key] = _parse_check_val(raw)
                    elif key == "isAbnormal":
                        rows_dict[row_idx][key] = safe_str(raw) in {"是", "✓", "true", "1"}
                    elif is_numeric_field_key(key):
                        rows_dict[row_idx][key] = safe_float(raw)
                    else:
                        rows_dict[row_idx][key] = safe_str(raw)
        wb.close()
        return list(rows_dict.values()), errors

    actual, raw = parse_upload_xlsx(content, _G11_5_FLAT_HEADERS, header_row=_HEADER_ROW)
    rows, truncated = import_rows_generic(
        raw, actual, _G11_5_FLAT_KEYS,
        parse_fn=lambda r, h: _parse_g11_5_row(r, h),
    )
    wb.close()
    if truncated:
        errors.append(f"数据行超过{ROW_LIMIT}行限制，已截断")
    return rows, errors


@router.post("/api/workpapers/{wp_id}/g11/export-template")
async def export_template(
    wp_id: str,
    sheet: str = Query(...),
    current_user: User = Depends(get_current_user),
) -> StreamingResponse:
    _validate(sheet)
    sp = _G11_SPECS[sheet]
    build_wb = sp.get("build_workbook")
    if callable(build_wb):
        wb = build_wb(None, template_only=True)
        return workbook_to_response(wb, f"{sheet}_模板.xlsx")
    if sheet == "G11-5":
        wb = _build_g11_5_workbook([], template_only=True)
        return workbook_to_response(wb, "G11-5_模板.xlsx")
    wb = build_workbook_template(
        sheet, sp["headers"], title=sp.get("title"), guidance=sp.get("guidance"),
    )
    return workbook_to_response(wb, f"{sheet}_模板.xlsx")


@router.post("/api/workpapers/{wp_id}/g11/export-data")
async def export_data(
    wp_id: str,
    sheet: str = Query(...),
    db: AsyncSession = Depends(get_db),
    current_user: User = Depends(get_current_user),
) -> StreamingResponse:
    _validate(sheet)
    sp = _G11_SPECS[sheet]
    build_wb = sp.get("build_workbook")
    if callable(build_wb):
        payload = await load_json_payload(db, wp_id, sp["item_id"], field=sp.get("storage_field", "remark"))
        wb = build_wb(payload, template_only=False)
        return workbook_to_response(wb, f"{sheet}_数据.xlsx")
    if sheet == "G11-1":
        store = await _load_g11_adj_remark(db, wp_id)
        wb = build_workbook_template(
            sheet, sp["headers"], title=sp.get("title"), guidance=sp.get("guidance"),
        )
        ws = wb[sheet]
        for row_vals in _export_g11_1_rows(store):
            ws.append(row_vals)
        return workbook_to_response(wb, f"{sheet}_数据.xlsx")
    if sheet == "G11-5":
        item_id = sp["item_id"]
        rows = await load_json_rows(db, wp_id, item_id, field="remark")
        wb = _build_g11_5_workbook(rows, template_only=False)
        return workbook_to_response(wb, "G11-5_数据.xlsx")
    item_id = sp["item_id"]
    rows = await load_json_rows(db, wp_id, item_id, field="remark")
    wb = build_workbook_template(
        sheet, sp["headers"], title=sp.get("title"), guidance=sp.get("guidance"),
    )
    ws = wb[sheet]
    for d in rows:
        ws.append(export_row_by_keys(d, sp["field_keys"]))
    return workbook_to_response(wb, f"{sheet}_数据.xlsx")


@router.post("/api/workpapers/{wp_id}/g11/import-data")
async def import_data(
    wp_id: str,
    sheet: str = Query(...),
    file: UploadFile = File(...),
    db: AsyncSession = Depends(get_db),
    current_user: User = Depends(get_current_user),
) -> dict[str, Any]:
    _validate(sheet)
    if not file.filename or not file.filename.endswith(".xlsx"):
        raise HTTPException(400, "请上传 .xlsx 格式文件")
    content = await file.read()
    if len(content) > 10 * 1024 * 1024:
        raise HTTPException(400, "文件大小不能超过10MB")

    sp = _G11_SPECS[sheet]
    parse_import = sp.get("parse_import")
    if callable(parse_import):
        payload, errors, imported_count = parse_import(content)
        if errors and not imported_count:
            return {"ok": False, "errors": errors, "imported_count": 0}
        await upsert_json_payload(
            db, wp_id, sp["item_id"], payload, field=sp.get("storage_field", "remark"),
        )
        return {"ok": True, "imported_count": imported_count, "errors": errors}

    if sheet == "G11-1":
        try:
            actual, raw = parse_upload_xlsx(content, sp["headers"], header_row=_HEADER_ROW)
        except ValueError as e:
            return {"ok": False, "errors": [str(e)], "imported_count": 0}
        parsed_rows: list[dict] = []
        for row_idx, row in enumerate(raw):
            if row_idx >= ROW_LIMIT:
                break
            parsed_rows.append(_parse_g11_1_row(row, actual))
        store = _rows_to_g11_1_store(parsed_rows)
        await _upsert_g11_adj_store(db, wp_id, store)
        return {"ok": True, "imported_count": len(store), "errors": []}

    if sheet == "G11-5":
        rows, errors = _parse_g11_5_upload(content)
        if errors and not rows:
            return {"ok": False, "errors": errors, "imported_count": 0}
        await upsert_json_rows(db, wp_id, sp["item_id"], rows, field="remark")
        return {"ok": True, "imported_count": len(rows), "errors": errors}

    try:
        actual, raw = parse_upload_xlsx(
            content,
            sp["headers"],
            header_row=_HEADER_ROW,
            header_aliases=sp.get("header_aliases"),
            require_all_headers=not bool(sp.get("allow_missing_headers")),
        )
    except ValueError as e:
        return {"ok": False, "errors": [str(e)], "imported_count": 0}
    keys = sp["field_keys"]
    expected = list(sp["headers"])
    rows, truncated = import_rows_generic(
        raw,
        actual,
        keys,
        parse_fn=lambda r, h: parse_row_by_headers(r, h, keys, expected_headers=expected),
    )
    # G11-3 旧值 AJE/RJE → 类别中文
    if sheet == "G11-3":
        for row in rows:
            cat = str(row.get("category") or "")
            if cat.upper() == "AJE":
                row["category"] = "账项调整"
            elif cat.upper() == "RJE":
                row["category"] = "报表调整"
            if not row.get("description") and row.get("summary"):
                row["description"] = row["summary"]
            if not row.get("indexRef") and row.get("index"):
                row["indexRef"] = row["index"]
    await upsert_json_rows(db, wp_id, sp["item_id"], rows, field="remark")
    out: dict[str, Any] = {"ok": True, "imported_count": len(rows), "errors": []}
    if truncated:
        out["warning"] = f"数据行数超过{ROW_LIMIT}行限制，已截断"
    return out

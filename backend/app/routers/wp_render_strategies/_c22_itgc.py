"""C22 IT 一般控制测试聚合 — 专属渲染策略.

component_type = "c22-itgc-bundle"

C22 为单一 34-sheet 工作簿（1 主矩阵 `C22 IT一般控制测试` + 33 子页），走 C1
多 sheet 整册模式（_WHOLE_WP_MULTISHEET_DEDICATED）。前端 GtC22ItgcBundle 为自加载
聚合组件：内部以「ITGC 矩阵总览 + SA/PE/PM/NS 分组页签 + C21/C21-1」切换，子页通过
GtOnlyOfficeSheet 按 sheetName 渲染同一工作簿的对应 sheet。

本渲染策略的关键作用（对齐 render_c1_entity_control）：
让 component_type 在 RENDERER_DISPATCH 中命中，避免多 sheet dispatch 循环把 C22
各 sheet 误判为非白名单而重写成 onlyoffice-sheet（铁律：D~N/专属组件必须注册）。

返回轻量 html_data（project_context + responses_snapshot），并在**主矩阵 sheet**
额外提取矩阵行内容（IT 控制类别 / 风险编号 / 控制编号 / 控制描述 / 相关应用系统 /
索引号），供前端 matrix 总览面板与子页公式引用消费（Phase0 §2 列结构）。

数据持久化在 checklist_responses 表，item_id 前缀为 "C22." （详见 phase0-notes §4/§5）。
"""

from __future__ import annotations

import logging
from pathlib import Path

import sqlalchemy as sa

from ._context import RenderContext

logger = logging.getLogger(__name__)

# ─── 主矩阵 sheet 名 + 列结构（Phase0 §2 实测：表头行=9，数据行 10–45） ─────────
MATRIX_SHEET_NAME = "C22 IT一般控制测试"
_MATRIX_FIRST_ROW = 10
_MATRIX_LAST_ROW = 45
# 列 → 字段（openpyxl 1-based 列号）
_COL_CATEGORY = 1      # A IT控制类别
_COL_RISK = 2          # B 风险描述（ITRxxx + 文字）
_COL_CONTROL_NO = 3    # C 控制编号（SA-3 / PE-3a / ...）
_COL_DESC = 4          # D 控制描述
_COL_APP_SYSTEM = 5    # E 相关应用系统
_COL_INDEX = 6         # F 索引号


def _extract_matrix_rows(template_path: str | None) -> list[dict]:
    """从模板主矩阵 sheet 提取控制点行内容（静态部分，动态结论走 responses）。

    仅提取 A~F 列（类别/风险/控制编号/描述/应用系统/索引），G~J（设计/执行结论/
    是否异常/缺陷编号）为审计填报字段，由前端从 checklist_responses 取。

    失败时返回 []（前端 matrix 面板降级为仅从 TabDef + responses 构建）。
    """
    if not template_path:
        return []
    p = Path(template_path)
    if not p.is_file() or p.suffix.lower() not in (".xlsx", ".xls"):
        return []
    try:
        import openpyxl

        wb = openpyxl.load_workbook(template_path, read_only=True, data_only=True)
        if MATRIX_SHEET_NAME not in wb.sheetnames:
            wb.close()
            return []
        ws = wb[MATRIX_SHEET_NAME]

        def _cell(row: int, col: int) -> str:
            try:
                v = ws.cell(row=row, column=col).value
            except Exception:  # noqa: BLE001
                return ""
            return str(v).strip() if v is not None else ""

        rows: list[dict] = []
        for r in range(_MATRIX_FIRST_ROW, _MATRIX_LAST_ROW + 1):
            category = _cell(r, _COL_CATEGORY)
            risk = _cell(r, _COL_RISK)
            control_no = _cell(r, _COL_CONTROL_NO)
            desc = _cell(r, _COL_DESC)
            app_system = _cell(r, _COL_APP_SYSTEM)
            index_no = _cell(r, _COL_INDEX)
            # 整行空则跳过（矩阵尾部空行）
            if not any((category, risk, control_no, desc, app_system, index_no)):
                continue
            rows.append(
                {
                    "row": r,
                    "category": category,
                    "risk": risk,
                    "controlNo": control_no,
                    "description": desc,
                    "appSystem": app_system,
                    "indexNo": index_no,
                }
            )
        wb.close()
        return rows
    except Exception as e:  # noqa: BLE001
        logger.warning("C22 render: 主矩阵提取失败 path=%s: %s", template_path, e)
        return []


async def render(ctx: RenderContext) -> dict | None:
    """C22 IT 一般控制测试渲染策略 — 返回轻量 html_data.

    返回 dict（非 grid cells），确保前端 GtWpRenderer 走 rendererEntry 分发到
    GtC22ItgcBundle，而非 grid 兜底或 OnlyOffice。

    - 主矩阵 sheet：额外返回 matrix 行内容（供 matrix 总览面板）。
    - 子页 sheet：仅返回 project_context + responses_snapshot（轻量）。
    """
    wp_id = ctx.wp_id
    db = ctx.db
    sheet_name = ctx.classification.sheet_name if ctx.classification else ""
    is_matrix = sheet_name == MATRIX_SHEET_NAME

    # ─── 从 checklist_responses 加载 C22.* 数据快照 ─────────────────────────
    responses_snapshot: dict = {}
    try:
        result = await db.execute(
            sa.text(
                "SELECT item_id, conclusion, remark "
                "FROM checklist_responses WHERE wp_id = :wp_id "
                "AND item_id LIKE 'C22.%' "
                "LIMIT 2000"
            ),
            {"wp_id": str(wp_id)},
        )
        for row in result.fetchall():
            responses_snapshot[row.item_id] = {
                "conclusion": row.conclusion or "",
                "remark": row.remark or "",
            }
    except Exception as e:  # noqa: BLE001
        logger.warning(
            "C22 render: checklist_responses 查询失败 wp_id=%s: %s", wp_id, e
        )

    # ─── 项目上下文 ────────────────────────────────────────────────────────
    project_context: dict = {
        "client_name": "",
        "audit_year": "",
        "business_category": ctx.business_category or "",
    }
    try:
        proj_row = (
            await db.execute(
                sa.text(
                    "SELECT client_name, audit_year, business_category "
                    "FROM projects WHERE id = :pid"
                ),
                {"pid": str(ctx.project_id)},
            )
        ).fetchone()
        if proj_row:
            project_context["client_name"] = proj_row.client_name or ""
            project_context["audit_year"] = str(proj_row.audit_year or "")
            project_context["business_category"] = (
                proj_row.business_category or ctx.business_category or ""
            )
    except Exception as e:  # noqa: BLE001
        logger.warning("C22 render: project context 查询失败: %s", e)

    out: dict = {
        "sheet_name": sheet_name,
        "is_matrix": is_matrix,
        "project_context": project_context,
        "responses_snapshot": responses_snapshot,
    }
    # 主矩阵 sheet 额外附带矩阵行内容（供总览面板）
    if is_matrix:
        out["matrix"] = _extract_matrix_rows(ctx.template_file_path)
    return out

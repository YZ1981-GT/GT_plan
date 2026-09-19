"""D2 应收账款「附注披露」导入导出（sheet=D2-disc-listed / D2-disc-soe）

由 `_d2_import_export.py` 的三个端点分派进来：
    POST /api/workpapers/{wp_id}/d2/export-template?sheet=D2-disc-listed
    POST /api/workpapers/{wp_id}/d2/export-data?sheet=D2-disc-soe
    POST /api/workpapers/{wp_id}/d2/import-data?sheet=D2-disc-listed

与其它 D2 sheet 的差异（故独立成模块，不塞进单 sheet 架构）：
  - 披露页由 ~10 张结构不同的表组成（列头逐字对齐附注模板 五、5 / 八、5），
    单个工作表放不下 → **多工作表工作簿**，一张表一个 worksheet + 一张「说明文本」。
  - 只导入导出**审计师手工录入**的明细表与说明文本；
    自动取数表（账龄披露、按计提方法分类、坏账准备变动）由 D2-2/D2-1/D2-3 取数派生，
    以及手工覆盖值（overrides）**不在导入导出范围**（避免用 Excel 绕过取数链与覆盖标记）。

存储键（与前端 `useD2DisclosureNote` 前缀一致，单一真源）：
    D2-disc-{variant}-individual-rows / -portfolios / -other-portfolio-rows
    D2-disc-{variant}-reversal-rows / -writeoff-rows / -top5-rows / -derecognized-rows
    D2-disc-{variant}-note-{子节key}（remark 存纯文本）
"""
from __future__ import annotations

import io
import json
from typing import Any
from uuid import uuid4

import sqlalchemy as sa
from fastapi import HTTPException
from fastapi.responses import StreamingResponse
from sqlalchemy.ext.asyncio import AsyncSession

from app.routers.wp_render_strategies._cycle_import_export_common import (
    load_json_rows,
    resolve_aging_segments,
    safe_float,
    safe_str,
    upsert_json_rows,
)

try:
    from openpyxl import Workbook
    from openpyxl.styles import Alignment, Font, PatternFill
except ImportError:  # pragma: no cover - 环境缺 openpyxl 时由上层端点报错
    Workbook = None  # type: ignore
    Font = Alignment = PatternFill = None  # type: ignore

# ─── Sheet 参数 ───────────────────────────────────────────────────────────────

DISCLOSURE_SHEETS: dict[str, str] = {
    'D2-disc-listed': 'listed',
    'D2-disc-soe': 'soe',
}

MAX_IMPORT_ROWS = 500

# ─── 说明文本子节（与前端 D2_NOTE_TEXT_SECTIONS 逐字一致）────────────────────

NOTE_SECTIONS: list[tuple[str, str]] = [
    ('aging', '按账龄披露说明'),
    ('badDebtClass', '按坏账准备计提方法分类说明'),
    ('individual', '按单项计提坏账准备说明'),
    ('portfolio', '按组合计提坏账准备说明'),
    ('movement', '坏账准备变动说明'),
    ('writeOff', '应收账款核销说明'),
    ('top5', '前五名欠款方说明'),
]

# ─── 工作表定义（表名 / 列头 / 字段 / 适用变体）───────────────────────────────
# field: (storage字段名, 是否数值)；组合计提项目为父子结构单独处理。

WS_INDIVIDUAL = '单项计提明细'
WS_PORTFOLIO = '组合计提项目'
WS_OTHER_PORTFOLIO = '其他组合方法'
WS_REVERSAL = '转回或收回'
WS_WRITEOFF = '核销逐项'
WS_TOP5 = '前五名'
WS_DERECOGNIZED = '终止确认'
WS_NOTES = '说明文本'

_INDIVIDUAL_FIELDS: list[tuple[str, str, bool]] = [
    ('name', '名称', False),
    ('endAmount', '期末余额', True),
    ('priorAmount', '上年年末余额', True),
    ('provision', '坏账准备', True),
    ('aging', '账龄', False),
    ('lossRate', '预期信用损失率(%)', True),
    ('basis', '计提理由', False),
]

_OTHER_PORTFOLIO_FIELDS: list[tuple[str, str, bool]] = [
    ('name', '组合名称', False),
    ('endAmount', '期末数', True),
    ('priorAmount', '期初数', True),
]

_REVERSAL_FIELDS: list[tuple[str, str, bool]] = [
    ('companyName', '单位名称', False),
    ('reversalReason', '转回原因', False),
    ('recoveryMethod', '收回方式', False),
    ('originalBasis', '原确定坏账准备的依据', False),
    ('cumulativeProvision', '转回或收回前累计已计提坏账准备金额', True),
    ('amount', '转回或收回金额', True),
]

_WRITEOFF_FIELDS: list[tuple[str, str, bool]] = [
    ('companyName', '单位名称', False),
    ('nature', '应收账款性质', False),
    ('amount', '核销金额', True),
    ('reason', '核销原因', False),
    ('procedure', '履行的核销程序', False),
    ('relatedParty', '是否由关联交易产生', False),
]

_TOP5_FIELDS: list[tuple[str, str, bool]] = [
    ('companyName', '单位名称', False),
    ('arAmount', '应收账款期末余额', True),
    ('contractAssetAmount', '合同资产期末余额', True),
    ('provision', '坏账准备', True),
]

_DERECOGNIZED_FIELDS: list[tuple[str, str, bool]] = [
    ('companyName', '债务人名称', False),
    ('amount', '终止确认金额', True),
    ('gainLoss', '与终止确认相关的利得或损失', True),
]

_PORTFOLIO_COLUMNS = ['组合名称', '账龄', '期末余额', '上年年末余额']

# 简单行表：worksheet 名 → (存储键后缀, 字段定义, 仅国企?)
_SIMPLE_TABLES: list[tuple[str, str, list[tuple[str, str, bool]], bool]] = [
    (WS_INDIVIDUAL, 'individual-rows', _INDIVIDUAL_FIELDS, False),
    (WS_OTHER_PORTFOLIO, 'other-portfolio-rows', _OTHER_PORTFOLIO_FIELDS, True),
    (WS_REVERSAL, 'reversal-rows', _REVERSAL_FIELDS, False),
    (WS_WRITEOFF, 'writeoff-rows', _WRITEOFF_FIELDS, False),
    (WS_TOP5, 'top5-rows', _TOP5_FIELDS, False),
    (WS_DERECOGNIZED, 'derecognized-rows', _DERECOGNIZED_FIELDS, True),
]


def _tables_for(variant: str) -> list[tuple[str, str, list[tuple[str, str, bool]]]]:
    return [
        (ws_name, suffix, fields)
        for ws_name, suffix, fields, soe_only in _SIMPLE_TABLES
        if (not soe_only) or variant == 'soe'
    ]


# ─── 通用小工具 ───────────────────────────────────────────────────────────────

def resolve_variant(sheet: str) -> str:
    variant = DISCLOSURE_SHEETS.get(sheet)
    if not variant:
        raise HTTPException(status_code=400, detail=f"不支持的披露sheet: {sheet}")
    return variant


def _prefix(variant: str) -> str:
    return f"D2-disc-{variant}-"


def _gen_row_id(prefix: str) -> str:
    return f"{prefix}-{uuid4().hex[:10]}"


def _normalize_label(text: Any) -> str:
    """账龄 label 归一（中文数字/全角括号/到-至 差异容错），用于导入匹配段。"""
    s = safe_str(text)
    if not s:
        return ''
    table = {
        '一': '1', '二': '2', '三': '3', '四': '4', '五': '5',
        '（': '(', '）': ')', '～': '-', '~': '-', '至': '-', '到': '-',
        '　': '', ' ': '',
    }
    return ''.join(table.get(ch, ch) for ch in s)


def _seg_key_by_label(segments: list[Any], label: Any) -> tuple[str, str] | None:
    target = _normalize_label(label)
    if not target:
        return None
    for seg in segments:
        seg_label = seg.get('label') if isinstance(seg, dict) else getattr(seg, 'label', '')
        seg_key = seg.get('key') if isinstance(seg, dict) else getattr(seg, 'key', '')
        if _normalize_label(seg_label) == target:
            return str(seg_key), str(seg_label)
    return None


# ─── 文本（说明）读写：remark 存纯文本，不能走 upsert_json_payload ─────────────

async def _load_texts(db: AsyncSession, wp_id: str, variant: str) -> dict[str, str]:
    prefix = _prefix(variant)
    rows = await db.execute(
        sa.text(
            "SELECT item_id, remark FROM checklist_responses "
            "WHERE wp_id = :wp_id AND item_id LIKE :pattern"
        ),
        {"wp_id": wp_id, "pattern": f"{prefix}note-%"},
    )
    out: dict[str, str] = {}
    for item_id, remark in rows.all():
        key = str(item_id)[len(prefix) + len('note-'):]
        out[key] = '' if remark is None else str(remark)
    return out


async def _upsert_text(db: AsyncSession, wp_id: str, item_id: str, text: str) -> None:
    proj = await db.execute(
        sa.text(
            "SELECT project_id FROM working_paper WHERE id = :wp_id AND is_deleted = false"
        ),
        {"wp_id": wp_id},
    )
    project_id = proj.scalar_one_or_none()
    if not project_id:
        raise HTTPException(status_code=404, detail=f"底稿不存在: {wp_id}")
    await db.execute(
        sa.text("""
            INSERT INTO checklist_responses (id, project_id, wp_id, item_id, remark, updated_at, created_at)
            VALUES (:id, :project_id, :wp_id, :item_id, :remark, NOW(), NOW())
            ON CONFLICT (wp_id, item_id)
            DO UPDATE SET remark = :remark, updated_at = NOW()
        """),
        {
            "id": str(uuid4()),
            "project_id": str(project_id),
            "wp_id": wp_id,
            "item_id": item_id,
            "remark": text,
        },
    )


# ─── 工作簿构建 ───────────────────────────────────────────────────────────────

def _style_header(ws, columns: list[str]) -> None:
    header_font = Font(bold=True, size=11)
    fill = PatternFill(start_color="DDEBF7", end_color="DDEBF7", fill_type="solid")
    for idx, name in enumerate(columns, start=1):
        cell = ws.cell(row=1, column=idx, value=name)
        cell.font = header_font
        cell.fill = fill
        cell.alignment = Alignment(horizontal="center", vertical="center", wrap_text=True)
        ws.column_dimensions[cell.column_letter].width = max(12, min(30, len(name) * 2 + 4))
    ws.freeze_panes = "A2"


def _build_instruction_ws(ws, variant: str, segments: list[Any]) -> None:
    bold = Font(bold=True)
    section = '五、5' if variant == 'listed' else '八、5'
    lines: list[tuple[str, bool]] = [
        (f"D2 应收账款附注披露（{'上市公司版 ' + section if variant == 'listed' else '国企版 ' + section}） — 编制说明", True),
        ("", False),
        ("一、范围", True),
        ("1. 本工作簿只覆盖披露页中**审计师手工录入**的明细表与各子节说明文本。", False),
        ("2. 以下表由取数派生，不在导入导出范围（请在底稿页核对/覆盖）：", False),
        ("   · 按账龄披露（取自 D2-2 明细表账龄）", False),
        ("   · 按坏账准备计提方法分类披露（取自 D2-1 审定表）", False),
        ("   · 坏账准备变动（取自 D2-3 坏账准备明细表）", False),
        ("   · 单元格手工覆盖值（overrides）不导出，导入亦不改动。", False),
        ("", False),
        ("二、工作表说明", True),
        (f"1. {WS_INDIVIDUAL}：单项计提坏账准备明细。上市版只用「名称/期末余额/上年年末余额」；", False),
        ("   国企版另用「坏账准备/账龄/预期信用损失率(%)/计提理由」。列全部保留以保证往返不丢字段。", False),
        (f"2. {WS_PORTFOLIO}：组合计提项目（父子结构展平）。同一「组合名称」的连续多行构成一张组合分表；", False),
        ("   「账龄」必须是当前项目账龄段之一，未匹配的行会被跳过并在返回中提示。", False),
        (f"3. {WS_REVERSAL} / {WS_WRITEOFF} / {WS_TOP5}：重要转回收回、核销逐项、前五名欠款方。", False),
        (f"4. {WS_NOTES}：各子节说明文本（子节名不可改，只改「说明」列）。", False),
    ]
    if variant == 'soe':
        lines.append((f"5. {WS_OTHER_PORTFOLIO} / {WS_DERECOGNIZED}：国企版专有表。", False))
    lines += [
        ("", False),
        ("三、导入规则", True),
        ("1. 导入按工作表覆盖对应存储（同一表整表替换），未出现的工作表不改动。", False),
        (f"2. 单表最多 {MAX_IMPORT_ROWS} 行，超出部分截断并返回警告。", False),
        ("3. 名称/单位名称全空且金额全为 0 的行视为空行跳过。", False),
        ("4. 某工作表无有效数据行（或说明为空）时**不清空**既有数据，只在返回中提示；", False),
        ("   如需清空某表，请在底稿页逐行删除。", False),
        ("", False),
        ("四、当前项目账龄段", True),
        ("   " + ("、".join(
            str(seg.get('label') if isinstance(seg, dict) else getattr(seg, 'label', ''))
            for seg in segments
        ) or "（未获取到账龄配置，请在项目设置 → 底稿配置中确认）"), False),
    ]
    row = 1
    for text, is_bold in lines:
        cell = ws.cell(row=row, column=1, value=text)
        if is_bold:
            cell.font = bold
        row += 1
    ws.column_dimensions['A'].width = 92


def build_disclosure_workbook(
    variant: str,
    segments: list[Any],
    data: dict[str, Any] | None = None,
) -> "Workbook":
    """构建披露导入导出工作簿（data 为 None 时是空白模板）。"""
    if Workbook is None:  # pragma: no cover
        raise HTTPException(status_code=500, detail="服务端缺少 openpyxl 依赖")

    wb = Workbook()
    ws_instr = wb.active
    ws_instr.title = '编制说明'
    _build_instruction_ws(ws_instr, variant, segments)

    stored = data or {}

    # 简单行表
    for ws_name, suffix, fields in _tables_for(variant):
        ws = wb.create_sheet(ws_name)
        _style_header(ws, [label for _, label, _ in fields])
        for row in stored.get(suffix, []) or []:
            ws.append([
                (safe_float(row.get(key)) if is_num else safe_str(row.get(key)))
                for key, _, is_num in fields
            ])

    # 组合计提项目（父子展平）
    ws_pf = wb.create_sheet(WS_PORTFOLIO)
    _style_header(ws_pf, _PORTFOLIO_COLUMNS)
    for group in stored.get('portfolios', []) or []:
        name = safe_str(group.get('name'))
        rows = group.get('rows') if isinstance(group, dict) else None
        for r in rows or []:
            ws_pf.append([
                name,
                safe_str(r.get('label')),
                safe_float(r.get('endAmount')),
                safe_float(r.get('priorAmount')),
            ])

    # 说明文本
    ws_notes = wb.create_sheet(WS_NOTES)
    _style_header(ws_notes, ['子节', '说明'])
    ws_notes.column_dimensions['B'].width = 80
    texts: dict[str, str] = stored.get('notes', {}) or {}
    for key, title in NOTE_SECTIONS:
        cell_text = texts.get(key, '')
        ws_notes.append([title, cell_text])
    for r in range(2, 2 + len(NOTE_SECTIONS)):
        ws_notes.cell(row=r, column=2).alignment = Alignment(wrap_text=True, vertical='top')

    return wb


# ─── 导出 ─────────────────────────────────────────────────────────────────────

async def _load_disclosure_data(db: AsyncSession, wp_id: str, variant: str) -> dict[str, Any]:
    prefix = _prefix(variant)
    out: dict[str, Any] = {}
    for _, suffix, _fields in _tables_for(variant):
        out[suffix] = await load_json_rows(db, wp_id, prefix + suffix, field='remark')
    # portfolios 是 [{groupId,name,rows:[...]}]，load_json_rows 对 list 原样返回
    out['portfolios'] = await load_json_rows(db, wp_id, prefix + 'portfolios', field='remark')
    out['notes'] = await _load_texts(db, wp_id, variant)
    return out


def _workbook_response(wb: "Workbook", filename: str) -> StreamingResponse:
    buf = io.BytesIO()
    wb.save(buf)
    buf.seek(0)
    from urllib.parse import quote as _quote

    fn = _quote(filename)
    return StreamingResponse(
        buf,
        media_type=(
            "application/vnd.openxmlformats-officedocument.spreadsheetml.sheet"
        ),
        headers={"Content-Disposition": f"attachment; filename*=UTF-8''{fn}"},
    )


async def export_disclosure_template(wp_id: str, sheet: str, db: AsyncSession) -> StreamingResponse:
    variant = resolve_variant(sheet)
    segments = await resolve_aging_segments(db, wp_id, 'D2')
    wb = build_disclosure_workbook(variant, segments, None)
    label = '上市公司' if variant == 'listed' else '国企'
    return _workbook_response(wb, f"D2应收账款附注披露({label})-模板.xlsx")


async def export_disclosure_data(wp_id: str, sheet: str, db: AsyncSession) -> StreamingResponse:
    variant = resolve_variant(sheet)
    segments = await resolve_aging_segments(db, wp_id, 'D2')
    data = await _load_disclosure_data(db, wp_id, variant)
    wb = build_disclosure_workbook(variant, segments, data)
    label = '上市公司' if variant == 'listed' else '国企'
    return _workbook_response(wb, f"D2应收账款附注披露({label})-数据.xlsx")


# ─── 导入 ─────────────────────────────────────────────────────────────────────

def _read_ws_rows(ws, columns: list[str]) -> tuple[list[dict[str, Any]], str | None]:
    """按首行表头读数据行（表头顺序容错，按列名匹配），返回 (rows, warning)。"""
    headers: list[str] = []
    for cell in ws[1]:
        headers.append(safe_str(cell.value))
    idx_by_col: dict[str, int] = {}
    for col in columns:
        if col in headers:
            idx_by_col[col] = headers.index(col)
    missing = [c for c in columns if c not in idx_by_col]
    rows: list[dict[str, Any]] = []
    warning: str | None = None
    if missing:
        return [], f"「{ws.title}」缺少列: {', '.join(missing)}"
    for row in ws.iter_rows(min_row=2, values_only=True):
        if row is None or all(v is None or safe_str(v) == '' for v in row):
            continue
        rows.append({col: (row[i] if i < len(row) else None) for col, i in idx_by_col.items()})
        if len(rows) >= MAX_IMPORT_ROWS:
            warning = f"「{ws.title}」超过 {MAX_IMPORT_ROWS} 行，已截断"
            break
    return rows, warning


def _row_is_empty(parsed: dict[str, Any], fields: list[tuple[str, str, bool]]) -> bool:
    for key, _label, is_num in fields:
        val = parsed.get(key)
        if is_num:
            if safe_float(val) != 0:
                return False
        elif safe_str(val):
            return False
    return True


def parse_disclosure_workbook(
    wb: Any, variant: str, segments: list[Any]
) -> tuple[dict[str, list[dict]], dict[str, str], list[str]]:
    """解析工作簿 → (存储后缀 → rows, 说明文本, 警告列表)。

    只处理工作簿中**存在**的工作表；未出现的表不写入（导入不误清空）。
    """
    storage: dict[str, list[dict]] = {}
    notes: dict[str, str] = {}
    warnings: list[str] = []

    for ws_name, suffix, fields in _tables_for(variant):
        if ws_name not in wb.sheetnames:
            continue
        raw_rows, warning = _read_ws_rows(wb[ws_name], [label for _, label, _ in fields])
        if warning and not raw_rows:
            warnings.append(warning)
            continue
        if warning:
            warnings.append(warning)
        parsed_rows: list[dict] = []
        for raw in raw_rows:
            parsed = {
                key: (safe_float(raw.get(label)) if is_num else safe_str(raw.get(label)))
                for key, label, is_num in fields
            }
            if _row_is_empty(parsed, fields):
                continue
            parsed['rowId'] = _gen_row_id(suffix[:3])
            parsed_rows.append(parsed)
        if not parsed_rows:
            # 空工作表（如直接导入空白模板）不清空既有数据，避免误删
            warnings.append(f"「{ws_name}」无有效数据行，未改动既有数据")
            continue
        storage[suffix] = parsed_rows

    # 组合计提项目（按「组合名称」分组，账龄需匹配当前段）
    if WS_PORTFOLIO in wb.sheetnames:
        raw_rows, warning = _read_ws_rows(wb[WS_PORTFOLIO], _PORTFOLIO_COLUMNS)
        if warning:
            warnings.append(warning)
        groups: list[dict[str, Any]] = []
        index_by_name: dict[str, int] = {}
        skipped_labels: list[str] = []
        for raw in raw_rows:
            name = safe_str(raw.get('组合名称'))
            matched = _seg_key_by_label(segments, raw.get('账龄'))
            if not name:
                continue
            if matched is None:
                label = safe_str(raw.get('账龄'))
                if label and label not in skipped_labels:
                    skipped_labels.append(label)
                continue
            seg_key, seg_label = matched
            if name not in index_by_name:
                index_by_name[name] = len(groups)
                groups.append({'groupId': _gen_row_id('pf'), 'name': name, 'rows': []})
            groups[index_by_name[name]]['rows'].append({
                'key': seg_key,
                'label': seg_label,
                'endAmount': safe_float(raw.get('期末余额')),
                'priorAmount': safe_float(raw.get('上年年末余额')),
            })
        if skipped_labels:
            warnings.append(
                f"「{WS_PORTFOLIO}」以下账龄未匹配当前账龄配置，已跳过: {', '.join(skipped_labels)}"
            )
        if groups:
            storage['portfolios'] = groups
        else:
            warnings.append(f"「{WS_PORTFOLIO}」无有效数据行，未改动既有数据")

    # 说明文本
    if WS_NOTES in wb.sheetnames:
        title_to_key = {title: key for key, title in NOTE_SECTIONS}
        for row in wb[WS_NOTES].iter_rows(min_row=2, values_only=True):
            if not row:
                continue
            title = safe_str(row[0] if len(row) > 0 else '')
            text = '' if len(row) < 2 or row[1] is None else str(row[1])
            key = title_to_key.get(title)
            if key is None:
                if title:
                    warnings.append(f"「{WS_NOTES}」未识别的子节名，已跳过: {title}")
                continue
            if not text.strip():
                # 空说明不覆盖既有文本（导入空白模板不清空说明）
                continue
            notes[key] = text

    return storage, notes, warnings


async def import_disclosure_data(wp_id: str, sheet: str, wb: Any, db: AsyncSession) -> dict[str, Any]:
    variant = resolve_variant(sheet)
    segments = await resolve_aging_segments(db, wp_id, 'D2')
    storage, notes, warnings = parse_disclosure_workbook(wb, variant, segments)

    if not storage and not notes:
        return {
            "data": {
                "rowCount": 0,
                "fieldCount": 0,
                "warning": "文件中无可识别的披露工作表",
                "warnings": warnings,
            }
        }

    prefix = _prefix(variant)
    row_count = 0
    for suffix, rows in storage.items():
        if suffix == 'portfolios':
            # portfolios 是父子结构（非扁平行），直接整体写入
            await _upsert_text(db, wp_id, prefix + suffix, json.dumps(rows, ensure_ascii=False))
            row_count += sum(len(g.get('rows') or []) for g in rows)
            continue
        await upsert_json_rows(db, wp_id, prefix + suffix, rows, field='remark')
        row_count += len(rows)

    for key, text in notes.items():
        await _upsert_text(db, wp_id, f"{prefix}note-{key}", text)

    await db.commit()

    result: dict[str, Any] = {
        "rowCount": row_count,
        "fieldCount": len(_INDIVIDUAL_FIELDS),
        "tables": {suffix: len(rows) for suffix, rows in storage.items()},
        "notes": len(notes),
    }
    if warnings:
        result["warnings"] = warnings
    return {"data": result}

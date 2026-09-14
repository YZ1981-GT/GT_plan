"""现金流量表核查 API 端点"""

from __future__ import annotations

from typing import Any
from uuid import UUID

from fastapi import APIRouter, Depends
from pydantic import BaseModel
from sqlalchemy.ext.asyncio import AsyncSession

from app.core.database import get_db
from app.services.cash_flow_verification_service import CashFlowVerificationService

router = APIRouter(prefix="/api/projects/{project_id}/cf-verification", tags=["cf-verification"])


class ExplanationRequest(BaseModel):
    check_type: str
    item_code: str | None = None
    explanation: str


@router.get("/{year}/cash-equivalents")
async def get_cash_equivalents(
    project_id: UUID,
    year: int,
    db: AsyncSession = Depends(get_db),
) -> dict[str, Any]:
    """现金及现金等价物列示 + BS 货币资金勾稽"""
    svc = CashFlowVerificationService(db)
    return await svc.get_cash_equivalents(project_id, year)


@router.get("/{year}/reconcile")
async def reconcile_bs_cf(
    project_id: UUID,
    year: int,
    db: AsyncSession = Depends(get_db),
) -> dict[str, Any]:
    """BS-CF 勾稽核对"""
    svc = CashFlowVerificationService(db)
    return await svc.reconcile_bs_cf(project_id, year)


@router.get("/{year}/main-table")
async def verify_main_table(
    project_id: UUID,
    year: int,
    db: AsyncSession = Depends(get_db),
) -> list[dict[str, Any]]:
    """主表逆算验证"""
    svc = CashFlowVerificationService(db)
    return await svc.verify_main_table(project_id, year)


@router.get("/{year}/supplementary")
async def verify_supplementary(
    project_id: UUID,
    year: int,
    db: AsyncSession = Depends(get_db),
) -> dict[str, Any]:
    """附表间接法核对"""
    svc = CashFlowVerificationService(db)
    return await svc.verify_supplementary(project_id, year)


@router.get("/{year}/results")
async def get_saved_results(
    project_id: UUID,
    year: int,
    db: AsyncSession = Depends(get_db),
) -> list[dict[str, Any]]:
    """获取已缓存的核查结果"""
    svc = CashFlowVerificationService(db)
    return await svc.get_saved_results(project_id, year)


@router.post("/{year}/explanation")
async def save_explanation(
    project_id: UUID,
    year: int,
    body: ExplanationRequest,
    db: AsyncSession = Depends(get_db),
) -> dict[str, Any]:
    """保存用户差异说明"""
    svc = CashFlowVerificationService(db)
    ok = await svc.save_explanation(project_id, year, body.check_type, body.item_code, body.explanation)
    if ok:
        await db.commit()
    return {"success": ok}


class CreateAdjustmentRequest(BaseModel):
    item_code: str
    difference: float
    description: str = ""


@router.post("/{year}/create-adjustment")
async def create_cf_adjustment(
    project_id: UUID,
    year: int,
    body: CreateAdjustmentRequest,
    db: AsyncSession = Depends(get_db),
) -> dict[str, Any]:
    """从 CF 差异一键创建调整分录"""
    from decimal import Decimal
    svc = CashFlowVerificationService(db)
    result = await svc.create_cf_adjustment_from_diff(
        project_id, year, body.item_code, Decimal(str(body.difference)), body.description
    )
    if result.get("success"):
        await db.commit()
    return result


@router.get("/{year}/export")
async def export_a5_excel(
    project_id: UUID,
    year: int,
    db: AsyncSession = Depends(get_db),
):
    """导出 A5-1 格式 Excel（现金流量表核查底稿）"""
    import io

    import openpyxl
    from fastapi.responses import StreamingResponse
    from openpyxl.styles import Alignment, Font, PatternFill

    svc = CashFlowVerificationService(db)

    # Gather data
    cash_data = await svc.get_cash_equivalents(project_id, year)
    recon_data = await svc.reconcile_bs_cf(project_id, year)
    main_data = await svc.verify_main_table(project_id, year)
    supp_data = await svc.verify_supplementary(project_id, year)

    wb = openpyxl.Workbook()
    header_font = Font(bold=True)
    header_fill = PatternFill(start_color="F4F0FA", end_color="F4F0FA", fill_type="solid")
    warn_font = Font(color="FF0000")

    # Sheet 1: 现金等价物
    ws = wb.active
    ws.title = "现金等价物"
    ws.append(["科目编码", "期初余额", "期末余额"])
    for cell in ws[1]:
        cell.font = header_font
        cell.fill = header_fill
    for item in cash_data.get("items", []):
        ws.append([item["account_code"], float(item["begin"]), float(item["end"])])
    ws.append([])
    ws.append(["TB合计", "", float(cash_data["total"])])
    ws.append(["BS货币资金", "", float(cash_data["bs_cash"])])
    ws.append(["差异", "", float(cash_data["difference"])])

    # Sheet 2: 勾稽核对
    ws2 = wb.create_sheet("勾稽核对")
    ws2.append(["项目", "金额"])
    for cell in ws2[1]:
        cell.font = header_font
        cell.fill = header_fill
    ws2.append(["BS货币资金期末", float(recon_data["bs_end"])])
    ws2.append(["BS货币资金期初", float(recon_data["bs_begin"])])
    ws2.append(["CF现金净增加额", float(recon_data["cf_net"])])
    ws2.append(["BS差异", float(recon_data["bs_diff"])])
    ws2.append([])
    ws2.append(["经营活动净额", float(recon_data["operating"])])
    ws2.append(["投资活动净额", float(recon_data["investing"])])
    ws2.append(["筹资活动净额", float(recon_data["financing"])])
    ws2.append(["汇率影响", float(recon_data["exchange"])])
    ws2.append(["三项合计", float(recon_data["activity_sum"])])
    ws2.append(["活动差异", float(recon_data["activity_diff"])])

    # Sheet 3: 主表逆算
    ws3 = wb.create_sheet("主表逆算")
    ws3.append(["活动", "项目", "行号", "报表值", "逆算值", "差异", "状态", "说明"])
    for cell in ws3[1]:
        cell.font = header_font
        cell.fill = header_fill
    section_map = {"operating": "经营", "investing": "投资", "financing": "筹资"}
    for item in main_data:
        row = [
            section_map.get(item["section"], item["section"]),
            item["item"],
            item["row_code"],
            float(item["reported"]) if item["reported"] is not None else "",
            float(item["calculated"]) if item["calculated"] is not None else "—",
            float(item["difference"]) if item["difference"] is not None else "—",
            "通过" if item["pass"] else ("人工" if item["calculated"] is None else "异常"),
            item.get("formula_description", ""),
        ]
        ws3.append(row)

    # Sheet 4: 附表间接法
    ws4 = wb.create_sheet("附表间接法")
    ws4.append(["调整项", "行号", "原始值", "符号", "调节值"])
    for cell in ws4[1]:
        cell.font = header_font
        cell.fill = header_fill
    for item in supp_data.get("items", []):
        ws4.append([
            item["item"],
            item["row_code"],
            float(item["value"]) if item["value"] is not None else 0,
            "+" if item["sign"] > 0 else "−",
            float(item["signed_value"]),
        ])
    ws4.append([])
    ws4.append(["间接法合计", "", float(supp_data["indirect_total"]), "", ""])
    ws4.append(["主表经营CF", "", float(supp_data["direct_operating"]), "", ""])
    ws4.append(["差异", "", float(supp_data["difference"]), "", ""])

    # Write to buffer
    buf = io.BytesIO()
    wb.save(buf)
    buf.seek(0)

    filename = f"A5-1_CF_Verification_{year}.xlsx"
    return StreamingResponse(
        buf,
        media_type="application/vnd.openxmlformats-officedocument.spreadsheetml.sheet",
        headers={"Content-Disposition": f"attachment; filename={filename}"},
    )

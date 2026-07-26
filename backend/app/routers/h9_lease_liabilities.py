"""H9 租赁负债 — 导入导出三级端点 + 摊销表生成API

4个端点：
- GET  /api/h9-lease-liabilities/{wp_id}/export-template  空白模板xlsx
- GET  /api/h9-lease-liabilities/{wp_id}/export-data      数据xlsx
- POST /api/h9-lease-liabilities/{wp_id}/import-data      解析xlsx写入
- POST /api/h9-lease-liabilities/{wp_id}/generate-amortization  摊销表生成（实际利率法）

支持sheets: H9-2/H9-3/H9-4
科目编码: 2205租赁负债（贷方/负债类）+ 未确认融资费用（借方/负债备抵类）
RFC5987 Content-Disposition header with Chinese filename encoding

Requirements: 3.6, 4.7
"""

from __future__ import annotations

import io
import json
import logging
from typing import Any
from urllib.parse import quote
from uuid import uuid4

from fastapi import APIRouter, Depends, File, HTTPException, Query, UploadFile
from fastapi.responses import StreamingResponse
from openpyxl import Workbook, load_workbook
from openpyxl.styles import Alignment, Font, PatternFill
from pydantic import BaseModel, Field
from sqlalchemy.ext.asyncio import AsyncSession

from app.core.database import get_db
from app.deps import get_current_user
from app.models.core import User

logger = logging.getLogger(__name__)

router = APIRouter(prefix="/api/h9-lease-liabilities", tags=["h9-lease-liabilities"])

# ═══════════════════════════════════════════════════════════════════════════════
# Sheet 配置
# ═══════════════════════════════════════════════════════════════════════════════

_ROW_LIMIT = 500

_SUPPORTED_SHEETS: set[str] = {"H9-2", "H9-3", "H9-4"}

# 列头与前端 useH9Detail / useH9FinanceCost / useH9Adjustment 字段对齐
_SHEET_HEADERS: dict[str, list[str]] = {
    "H9-2": [
        "出租方", "合同号", "承租资产", "IBR利率(%)", "租赁期(月)",
        "期初余额", "本期偿还", "本期利息",
        "期初AJE", "偿还AJE", "利息AJE", "重分类(一年内到期)",
        "1年以内", "1-2年", "2-3年", "3年以上",
        "是否关联方", "是否发函", "备注",
    ],
    "H9-3": [
        "出租方", "合同号", "期初余额", "本期增加(借)", "本期确认(贷)",
        "期初AJE", "增加AJE", "确认AJE", "其他AJE", "重分类",
        "1年以内", "1-2年", "2-3年", "3年以上",
        "对应利息期", "是否关联方", "备注",
    ],
    "H9-4": [
        "调整事项说明", "类别", "报表项目", "科目名称", "附注项目",
        "借方金额", "贷方金额", "索引", "备注",
    ],
}

# 中文列名 → JSON 字段名（与前端 composable 一致）
_FIELD_MAPS: dict[str, dict[str, str]] = {
    "H9-2": {
        "出租方": "lessor",
        "合同号": "contractNo",
        "承租资产": "assetDesc",
        "IBR利率(%)": "ibrRate",
        "租赁期(月)": "leaseTerm",
        "期初余额": "beginBalance",
        "本期偿还": "repayment",
        "本期利息": "interestAccrued",
        "期初AJE": "beginAje",
        "偿还AJE": "repayAje",
        "利息AJE": "interestAje",
        "重分类(一年内到期)": "reclassification",
        "1年以内": "dueWithin1Y",
        "1-2年": "due1To2Y",
        "2-3年": "due2To3Y",
        "3年以上": "dueOver3Y",
        "是否关联方": "isRelatedParty",
        "是否发函": "isConfirmed",
        "备注": "remark",
    },
    "H9-3": {
        "出租方": "lessor",
        "合同号": "contractNo",
        "期初余额": "beginBalance",
        "本期增加(借)": "debitIncrease",
        "本期确认(贷)": "creditDecrease",
        "期初AJE": "beginAje",
        "增加AJE": "increaseAje",
        "确认AJE": "confirmAje",
        "其他AJE": "otherAje",
        "重分类": "reclassification",
        "1年以内": "dueWithin1Y",
        "1-2年": "due1To2Y",
        "2-3年": "due2To3Y",
        "3年以上": "dueOver3Y",
        "对应利息期": "interestPeriod",
        "是否关联方": "isRelatedParty",
        "备注": "remark",
    },
    "H9-4": {
        "调整事项说明": "description",
        "类别": "category",
        "报表项目": "reportItem",
        "科目名称": "accountName",
        "附注项目": "noteItem",
        "借方金额": "debitAmount",
        "贷方金额": "creditAmount",
        "索引": "indexRef",
        "备注": "remark",
    },
}

_NUMERIC_FIELDS: dict[str, set[str]] = {
    "H9-2": {
        "ibrRate", "leaseTerm", "beginBalance", "repayment", "interestAccrued",
        "beginAje", "repayAje", "interestAje", "reclassification",
        "dueWithin1Y", "due1To2Y", "due2To3Y", "dueOver3Y",
    },
    "H9-3": {
        "beginBalance", "debitIncrease", "creditDecrease",
        "beginAje", "increaseAje", "confirmAje", "otherAje", "reclassification",
        "dueWithin1Y", "due1To2Y", "due2To3Y", "dueOver3Y",
    },
    "H9-4": {"debitAmount", "creditAmount"},
}

# checklist_responses item_id（与前端 useH9* 一致）
_SHEET_ITEM_ID: dict[str, str] = {
    "H9-2": "H9-2-rows",
    "H9-3": "H9-3-rows",
    "H9-4": "H9-4-rows",
}

# 历史/错误键别名（导出时兜底读取）
_SHEET_ITEM_ID_ALIASES: dict[str, list[str]] = {
    "H9-2": ["H9-detail-rows"],
    "H9-3": ["H9-finance-cost-rows"],
    "H9-4": ["H9-5-rows", "H9-adjustment-rows"],
}


# ═══════════════════════════════════════════════════════════════════════════════
# Pydantic 模型
# ═══════════════════════════════════════════════════════════════════════════════


class AmortizationRequest(BaseModel):
    """摊销表生成请求"""
    initial_balance: float = Field(..., gt=0, description="初始余额（租赁负债初始确认金额）")
    payment: float = Field(..., gt=0, description="每期付款金额")
    rate: float = Field(..., ge=0, lt=1, description="每期利率（非年化，如月利率=年利率/12）")
    periods: int = Field(..., gt=0, le=360, description="总期数")


class AmortizationRow(BaseModel):
    """摊销表单行"""
    period: int
    begin_balance: float
    payment: float
    interest: float
    principal: float
    end_balance: float


class AmortizationResponse(BaseModel):
    """摊销表生成响应"""
    schedule: list[AmortizationRow]
    validation: dict


# ═══════════════════════════════════════════════════════════════════════════════
# 辅助函数
# ═══════════════════════════════════════════════════════════════════════════════


def _validate_sheet(sheet_code: str) -> None:
    if sheet_code not in _SUPPORTED_SHEETS:
        raise HTTPException(
            400,
            f"不支持的sheet: {sheet_code}。支持: {sorted(_SUPPORTED_SHEETS)}",
        )


def _safe_float(val: Any) -> float:
    if val is None:
        return 0.0
    try:
        return float(val)
    except (ValueError, TypeError):
        return 0.0


def _safe_str(val: Any) -> str:
    if val is None:
        return ""
    return str(val).strip()


def _create_template_wb(sheet_code: str) -> Workbook:
    """创建空白模板xlsx（含表头+格式，无数据行）"""
    wb = Workbook()
    ws = wb.active
    ws.title = sheet_code

    headers = _SHEET_HEADERS[sheet_code]
    header_font = Font(bold=True, size=11)
    header_fill = PatternFill(start_color="D9E1F2", end_color="D9E1F2", fill_type="solid")
    header_align = Alignment(horizontal="center", vertical="center", wrap_text=True)

    for col_idx, col_name in enumerate(headers, 1):
        cell = ws.cell(row=1, column=col_idx, value=col_name)
        cell.font = header_font
        cell.fill = header_fill
        cell.alignment = header_align
        ws.column_dimensions[cell.column_letter].width = max(len(col_name) * 2 + 4, 12)

    ws.freeze_panes = "A2"
    return wb


def _validate_columns(ws: Any, sheet_code: str) -> list[str]:
    """校验列头，返回不匹配的列名列表"""
    expected = set(_SHEET_HEADERS[sheet_code])
    actual: list[str] = []
    for cell in next(ws.iter_rows(min_row=1, max_row=1)):
        if cell.value is not None:
            actual.append(str(cell.value).strip())

    missing = [h for h in expected if h not in actual]
    return missing


def _get_actual_headers(ws: Any) -> list[str]:
    """获取worksheet实际列头"""
    headers: list[str] = []
    for cell in next(ws.iter_rows(min_row=1, max_row=1)):
        if cell.value is not None:
            headers.append(str(cell.value).strip())
    return headers


def _col_val(row: tuple, actual_headers: list[str], col_name: str) -> Any:
    """从行数据中按列名取值"""
    try:
        idx = actual_headers.index(col_name)
        return row[idx] if idx < len(row) else None
    except (ValueError, IndexError):
        return None


def _export_row(sheet_code: str, data: dict) -> list:
    """通用导出行转换——根据sheet_code按列头顺序提取字段"""
    headers = _SHEET_HEADERS[sheet_code]
    field_map = _FIELD_MAPS.get(sheet_code, {})
    result = []
    for col_name in headers:
        field = field_map.get(col_name, col_name)
        val = data.get(field, "")
        # H9-4 类别导出为中文，便于 Excel 模板阅读
        if sheet_code == "H9-4" and field == "category":
            cat = str(val or "").upper()
            if cat in ("AJE", "账项调整"):
                val = "账项调整"
            elif cat in ("RJE", "报表调整"):
                val = "报表调整"
            else:
                val = val or "账项调整"
        result.append(val if val is not None else "")
    return result


def _parse_row(sheet_code: str, row: tuple, actual_headers: list[str]) -> dict:
    """通用导入行解析——根据sheet_code按列名映射回JSON字段"""
    field_map = _FIELD_MAPS.get(sheet_code, {})
    result: dict[str, Any] = {"rowId": str(uuid4())}
    for col_name, field_name in field_map.items():
        raw = _col_val(row, actual_headers, col_name)
        if field_name in _NUMERIC_FIELDS.get(sheet_code, set()):
            result[field_name] = _safe_float(raw)
        else:
            result[field_name] = _safe_str(raw)

    # H9-4 类别归一：账项调整/AJE → AJE；报表调整/RJE → RJE
    if sheet_code == "H9-4":
        cat = str(result.get("category") or "").strip().upper()
        if "报表" in cat or cat == "RJE":
            result["category"] = "RJE"
        else:
            result["category"] = "AJE"
        result.setdefault("seq", 0)

    # H9-2/H9-3 关联方默认
    if sheet_code in ("H9-2", "H9-3"):
        if not result.get("isRelatedParty"):
            result["isRelatedParty"] = "否"
    if sheet_code == "H9-2" and not result.get("isConfirmed"):
        result["isConfirmed"] = "否"

    return result


async def _load_sheet_rows(db: AsyncSession, wp_id: str, sheet: str) -> list[dict]:
    """按主键再别名读取 checklist 行数据"""
    import sqlalchemy as sa

    item_ids = [_SHEET_ITEM_ID[sheet], *_SHEET_ITEM_ID_ALIASES.get(sheet, [])]
    for item_id in item_ids:
        result = await db.execute(
            sa.text(
                "SELECT remark FROM checklist_responses "
                "WHERE wp_id = :wp_id AND item_id = :item_id LIMIT 1"
            ),
            {"wp_id": wp_id, "item_id": item_id},
        )
        row = result.fetchone()
        if not row or not row[0]:
            continue
        try:
            data = json.loads(row[0])
            if isinstance(data, list):
                return data
        except (json.JSONDecodeError, TypeError):
            continue
    return []


def _generate_amortization_schedule(
    initial_balance: float,
    payment: float,
    rate: float,
    periods: int,
) -> list[dict]:
    """生成实际利率法摊销表

    核心逻辑：
    - 每期利息 = 期初余额 × 利率
    - 本金偿还 = 每期付款 - 利息
    - 期末余额 = 期初 - 本金偿还
    - 最后一期调整尾差确保期末=0
    """
    schedule: list[dict] = []
    balance = initial_balance

    for period in range(1, periods + 1):
        interest = round(balance * rate, 2)

        if period == periods:
            # 最后一期：调整尾差，确保期末余额=0
            principal = balance
            actual_payment = principal + interest
            end_balance = 0.0
        else:
            principal = round(payment - interest, 2)
            end_balance = round(balance - principal, 2)
            actual_payment = payment

        schedule.append({
            "period": period,
            "begin_balance": round(balance, 2),
            "payment": round(actual_payment, 2),
            "interest": interest,
            "principal": round(principal, 2),
            "end_balance": end_balance,
        })

        balance = end_balance

    return schedule


# ═══════════════════════════════════════════════════════════════════════════════
# 端点
# ═══════════════════════════════════════════════════════════════════════════════


@router.get("/{wp_id}/export-template")
async def h9_export_template(
    wp_id: str,
    sheet: str = Query(..., description="Sheet编码: H9-2/H9-3/H9-4"),
    current_user: User = Depends(get_current_user),
) -> StreamingResponse:
    """导出空白xlsx模板（含表头+格式，无数据行）

    Requirements: 3.6
    """
    _validate_sheet(sheet)

    wb = _create_template_wb(sheet)
    buffer = io.BytesIO()
    wb.save(buffer)
    buffer.seek(0)

    filename = f"H9租赁负债_{sheet}_模板.xlsx"
    return StreamingResponse(
        buffer,
        media_type="application/vnd.openxmlformats-officedocument.spreadsheetml.sheet",
        headers={"Content-Disposition": f"attachment; filename*=UTF-8''{quote(filename)}"},
    )


@router.get("/{wp_id}/export-data")
async def h9_export_data(
    wp_id: str,
    sheet: str = Query(..., description="Sheet编码: H9-2/H9-3/H9-4"),
    db: AsyncSession = Depends(get_db),
    current_user: User = Depends(get_current_user),
) -> StreamingResponse:
    """导出含当前数据的xlsx

    Requirements: 3.6
    """
    _validate_sheet(sheet)
    rows_data = await _load_sheet_rows(db, wp_id, sheet)

    # 生成 xlsx
    wb = _create_template_wb(sheet)
    ws = wb.active

    for data_row in rows_data:
        ws.append(_export_row(sheet, data_row))

    buffer = io.BytesIO()
    wb.save(buffer)
    buffer.seek(0)

    filename = f"H9租赁负债_{sheet}_数据.xlsx"
    return StreamingResponse(
        buffer,
        media_type="application/vnd.openxmlformats-officedocument.spreadsheetml.sheet",
        headers={"Content-Disposition": f"attachment; filename*=UTF-8''{quote(filename)}"},
    )


@router.post("/{wp_id}/import-data")
async def h9_import_data(
    wp_id: str,
    sheet: str = Query(..., description="Sheet编码: H9-2/H9-3/H9-4"),
    file: UploadFile = File(...),
    db: AsyncSession = Depends(get_db),
    current_user: User = Depends(get_current_user),
) -> dict:
    """导入xlsx解析写入checklist_responses

    Requirements: 3.6
    """
    _validate_sheet(sheet)
    import sqlalchemy as sa

    # 读取上传文件
    content = await file.read()
    if len(content) > 10 * 1024 * 1024:
        raise HTTPException(400, "文件大小超过 10MB 限制")

    try:
        wb = load_workbook(io.BytesIO(content), read_only=True, data_only=True)
    except Exception:
        raise HTTPException(400, "无法解析xlsx文件")

    ws = wb.active

    # 校验列头
    missing = _validate_columns(ws, sheet)
    if missing:
        raise HTTPException(400, f"列名不匹配，缺少: {missing}")

    actual_headers = _get_actual_headers(ws)

    # 解析数据行
    parsed_rows: list[dict] = []
    for row_idx, row in enumerate(ws.iter_rows(min_row=2, values_only=True), start=2):
        if row_idx > _ROW_LIMIT + 1:
            break
        # 跳过全空行
        if all(cell is None or str(cell).strip() == "" for cell in row):
            continue
        parsed_rows.append(_parse_row(sheet, row, actual_headers))

    # 写入 checklist_responses
    item_id = _SHEET_ITEM_ID[sheet]
    json_str = json.dumps(parsed_rows, ensure_ascii=False)

    # 查 project_id（checklist_responses.project_id 为 NOT NULL）
    pid_result = await db.execute(
        sa.text("SELECT project_id FROM working_paper WHERE id = :wp_id LIMIT 1"),
        {"wp_id": wp_id},
    )
    pid_row = pid_result.fetchone()
    if not pid_row:
        raise HTTPException(404, "底稿不存在")
    project_id = str(pid_row[0])

    await db.execute(
        sa.text(
            "INSERT INTO checklist_responses (id, wp_id, project_id, item_id, remark) "
            "VALUES (:id, :wp_id, :project_id, :item_id, :remark) "
            "ON CONFLICT (wp_id, item_id) DO UPDATE SET remark = :remark"
        ),
        {
            "id": str(uuid4()),
            "wp_id": wp_id,
            "project_id": project_id,
            "item_id": item_id,
            "remark": json_str,
        },
    )
    await db.commit()

    return {"row_count": len(parsed_rows), "sheet": sheet}


@router.post("/{wp_id}/generate-amortization")
async def h9_generate_amortization(
    wp_id: str,
    body: AmortizationRequest,
    current_user: User = Depends(get_current_user),
) -> dict:
    """生成租赁负债摊销表（实际利率法）

    核心CAS21计算：
    - 每期利息 = 期初余额 × 实际利率
    - 本金偿还 = 每期付款 - 利息费用
    - 期末余额 = 期初余额 - 本金偿还
    - 最后一期调整尾差确保期末余额=0

    Requirements: 4.7
    """
    schedule = _generate_amortization_schedule(
        initial_balance=body.initial_balance,
        payment=body.payment,
        rate=body.rate,
        periods=body.periods,
    )

    # 验证：最后一期期末余额应≈0
    tail_diff = abs(schedule[-1]["end_balance"]) if schedule else 0.0
    is_valid = tail_diff < 1.0  # 允许±1元尾差

    return {
        "schedule": schedule,
        "validation": {
            "is_valid": is_valid,
            "tail_diff": tail_diff,
            "total_interest": round(sum(row["interest"] for row in schedule), 2),
            "total_principal": round(sum(row["principal"] for row in schedule), 2),
        },
    }

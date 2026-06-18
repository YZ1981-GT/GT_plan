"""Word 导出引擎（PRE-2，完成阶段公共基础设施）.

将结构化内容渲染进致同标准 Word 模板，统一实现"颜色语义"处理：

| 类型 | 识别                       | export 行为              |
|------|----------------------------|--------------------------|
| 红   | 公司名/年度/XX/201X        | 替换为项目数据（替换后转黑） |
| 蓝   | 【】编制提示               | **删除**（含 run）        |
| 黑   | 固定正文                   | 保留                     |
| 注释表 | 模板末尾参考/注释表       | **删除**（整张 table）    |

设计要点（见 .kiro/specs/completion-phase-infra/requirements.md PRE-2）：

- 本模块是**新建独立模块**，与已废弃的 ``word_template_filler.py``
  （``WordTemplateFiller``，deprecated）不同。颜色语义只在本 filler 实现，
  **禁止**在编排层（a17_word_exporter / regulatory_letter_service 等）重复实现。
- 颜色判定基于 python-docx ``run.font.color.rgb``：
  蓝 = ``red < 100 且 blue > 150``；红 = ``red > 200 且 green < 100``。
- 注释表删除采用**可配置关键词策略**（按表首单元格关键词识别，不写死）。
- 占位符替换同时覆盖 paragraphs 与 tables。
- python-docx 不可用时优雅降级（参考 ``download_template_prefilled`` 的 ImportError 分支）。
- 不依赖 DB（纯 docx 处理）；编排服务负责取数后传 ``context``。
"""

from __future__ import annotations

import io
import logging
from dataclasses import dataclass, field
from typing import Any

logger = logging.getLogger(__name__)

# 注释表识别默认关键词：表首单元格（或任一前置单元格）含以下任一关键词即视为注释/参考表删除。
# 可配置：调用方可经 ColorSemanticsConfig.note_table_keywords 覆盖，不写死。
DEFAULT_NOTE_TABLE_KEYWORDS: tuple[str, ...] = (
    "注：",
    "注:",
    "参考",
    "说明",
    "填表说明",
    "编制说明",
    "备注：",
)

# 蓝色编制提示常见包裹符号（用于辅助识别，仅作日志/调试用途；删除以颜色为准）。
GUIDANCE_BRACKETS: tuple[tuple[str, str], ...] = (("【", "】"),)


@dataclass
class ColorSemanticsConfig:
    """颜色语义 / 注释表删除策略（可配置）.

    Attributes:
        note_table_keywords: 注释表识别关键词（表首单元格命中即删除整表）。
        delete_blue_runs: 是否删除蓝色 run（编制提示）。
        recolor_red_to_black: 红色 run 占位符替换后是否转黑。
        delete_note_tables: 是否删除末尾注释/参考表。
        scan_note_table_cells: 扫描表格前 N 个单元格寻找关键词（默认仅首单元格=1）。
    """

    note_table_keywords: tuple[str, ...] = DEFAULT_NOTE_TABLE_KEYWORDS
    delete_blue_runs: bool = True
    recolor_red_to_black: bool = True
    delete_note_tables: bool = True
    scan_note_table_cells: int = 1


# ─── 颜色判定 ────────────────────────────────────────────────────────────────


def _run_color_rgb(run: Any):
    """安全读取 run 的字体颜色 RGB；无颜色返回 None."""
    try:
        color = run.font.color
        if color is not None and color.rgb is not None:
            return color.rgb
    except (AttributeError, ValueError):
        return None
    return None


def is_blue_run(run: Any) -> bool:
    """蓝色编制提示判定：red < 100 且 blue > 150."""
    rgb = _run_color_rgb(run)
    if rgb is None:
        return False
    return rgb[0] < 100 and rgb[2] > 150


def is_red_run(run: Any) -> bool:
    """红色项目数据占位符判定：red > 200 且 green < 100."""
    rgb = _run_color_rgb(run)
    if rgb is None:
        return False
    return rgb[0] > 200 and rgb[1] < 100


# ─── 占位符替换 ──────────────────────────────────────────────────────────────


def build_replacements(context: dict[str, Any]) -> dict[str, str]:
    """从项目 context 构造占位符替换映射.

    复用 ``download_template_prefilled`` 的映射约定：
    - ××公司 / ABC公司 / XX公司 → client_name
    - 202X → audit_year
    - 201X → 上年度（audit_year - 1）

    context 还可直接携带额外的 ``placeholders`` dict（精确占位符 → 值），
    会合并进结果（优先级高于约定映射）。
    """
    client_name = str(context.get("client_name") or "").strip() or "XX公司"
    audit_year_raw = context.get("audit_year")
    audit_year = str(audit_year_raw).strip() if audit_year_raw not in (None, "") else "202X"

    prev_year = context.get("prev_year")
    if prev_year not in (None, ""):
        prev_year_str = str(prev_year).strip()
    elif audit_year.isdigit():
        prev_year_str = str(int(audit_year) - 1)
    else:
        prev_year_str = "201X"

    replacements: dict[str, str] = {
        "××公司": client_name,
        "ABC公司": client_name,
        "XX公司": client_name,
        "202X": audit_year,
        "201X": prev_year_str,
    }

    extra = context.get("placeholders")
    if isinstance(extra, dict):
        for k, v in extra.items():
            if k:
                replacements[str(k)] = "" if v is None else str(v)

    return replacements


def _replace_in_run(run: Any, replacements: dict[str, str]) -> bool:
    """在单个 run 文本内执行占位符替换，返回是否有替换发生."""
    text = run.text
    if not text:
        return False
    new_text = text
    for key, val in replacements.items():
        if key in new_text:
            new_text = new_text.replace(key, val)
    if new_text != text:
        run.text = new_text
        return True
    return False


# ─── 段落级处理 ──────────────────────────────────────────────────────────────


def _process_paragraph(paragraph: Any, replacements: dict[str, str], config: ColorSemanticsConfig) -> None:
    """处理单个段落的 runs：蓝删、红替换+转黑、黑保留."""
    try:
        from docx.shared import RGBColor
    except ImportError:  # pragma: no cover - 已在入口校验
        return

    runs_to_remove = []
    for run in paragraph.runs:
        if config.delete_blue_runs and is_blue_run(run):
            runs_to_remove.append(run)
            continue
        if is_red_run(run):
            # 红色 run：替换占位符 → 转黑
            _replace_in_run(run, replacements)
            if config.recolor_red_to_black:
                try:
                    run.font.color.rgb = RGBColor(0, 0, 0)
                except (AttributeError, ValueError):
                    pass
        else:
            # 黑色 / 无色正文：保留文本，但仍执行占位符替换（模板正文里也可能有占位符）
            _replace_in_run(run, replacements)

    for run in runs_to_remove:
        _remove_run(run)


def _remove_run(run: Any) -> None:
    """从 XML 树中移除一个 run."""
    element = run._r
    parent = element.getparent()
    if parent is not None:
        parent.remove(element)


# ─── 注释表识别与删除 ────────────────────────────────────────────────────────


def _table_lead_text(table: Any, scan_cells: int) -> str:
    """提取表格前若干单元格的拼接文本（用于注释表关键词匹配）."""
    texts: list[str] = []
    count = 0
    for row in table.rows:
        for cell in row.cells:
            texts.append(cell.text or "")
            count += 1
            if count >= scan_cells:
                return " ".join(texts)
    return " ".join(texts)


def is_note_table(table: Any, config: ColorSemanticsConfig) -> bool:
    """注释/参考表判定：表首单元格（前 N 格）命中任一关键词."""
    if not config.note_table_keywords:
        return False
    lead = _table_lead_text(table, max(1, config.scan_note_table_cells))
    return any(kw in lead for kw in config.note_table_keywords)


def _remove_table(table: Any) -> None:
    """从文档中删除一张表."""
    element = table._tbl
    parent = element.getparent()
    if parent is not None:
        parent.remove(element)


def _process_tables(doc: Any, replacements: dict[str, str], config: ColorSemanticsConfig) -> int:
    """处理文档内所有表格：注释表删除 + 单元格内占位符/颜色语义处理.

    返回删除的注释表数量。
    """
    deleted = 0
    # 先收集需删除的注释表（避免迭代时修改）
    tables = list(doc.tables)
    for table in tables:
        if config.delete_note_tables and is_note_table(table, config):
            _remove_table(table)
            deleted += 1
            continue
        # 非注释表：处理单元格内段落
        for row in table.rows:
            for cell in row.cells:
                for paragraph in cell.paragraphs:
                    _process_paragraph(paragraph, replacements, config)
    return deleted


# ─── 公共入口 ────────────────────────────────────────────────────────────────


@dataclass
class FillResult:
    """填充结果元数据（供测试 / 编排服务断言）."""

    note_tables_deleted: int = 0
    degraded: bool = False
    extras: dict[str, Any] = field(default_factory=dict)


def _docx_available() -> bool:
    try:
        import docx  # noqa: F401
        return True
    except ImportError:
        return False


def fill_and_export(
    doc: Any,
    context: dict[str, Any] | None = None,
    config: ColorSemanticsConfig | None = None,
) -> tuple[Any, FillResult]:
    """对一个已打开的 python-docx ``Document`` 应用颜色语义 + 占位符 + 注释表删除.

    Args:
        doc: python-docx ``Document`` 对象（in-place 修改）。
        context: 项目数据 dict（client_name / audit_year / prev_year / placeholders）。
        config: 颜色语义 / 注释表删除策略；None 用默认。

    Returns:
        (doc, FillResult) — doc 为同一对象（in-place）；FillResult 含删除注释表数等元数据。

    Note:
        python-docx 不可用时优雅降级：原样返回 doc，FillResult.degraded=True。
    """
    context = context or {}
    config = config or ColorSemanticsConfig()

    if not _docx_available():
        logger.warning("python-docx 不可用，跳过颜色语义处理（优雅降级）")
        return doc, FillResult(degraded=True)

    replacements = build_replacements(context)

    # 1. 顶层段落（红替换转黑 / 蓝删 / 黑保留 + 占位符）
    for paragraph in doc.paragraphs:
        _process_paragraph(paragraph, replacements, config)

    # 2. 表格：注释表删除 + 单元格内处理
    deleted = _process_tables(doc, replacements, config)

    return doc, FillResult(note_tables_deleted=deleted)


def fill_template_path(
    template_path: str,
    context: dict[str, Any] | None = None,
    config: ColorSemanticsConfig | None = None,
) -> tuple[Any, FillResult]:
    """打开模板文件路径并应用颜色语义处理.

    python-docx 不可用时抛 ImportError 由调用方降级处理（与端点层 FileResponse 一致）。
    """
    from docx import Document

    doc = Document(template_path)
    return fill_and_export(doc, context, config)


def export_to_bytes(doc: Any) -> bytes:
    """将处理后的 ``Document`` 序列化为 docx 字节流（供 HTTP 响应 / 落盘）."""
    buffer = io.BytesIO()
    doc.save(buffer)
    return buffer.getvalue()

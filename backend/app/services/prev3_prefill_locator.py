"""三参 ``PREV('底稿','表','字段')`` 预填落格：按模板逐格的定位表（spec formula-fetch-followup-defect-closure R2）。

映射表 ``prefill_formula_mapping.json`` 里三参 PREV 格的位置是语义名（「上年审定数」等）。旧规则按列头关键词猜格
（列头含「上年」或「审定数」+ 第 7 行起第一个数据行），在模板上落到「期末与期初审定数比较」列或分节标题行。
定位表 ``backend/data/prev3_prefill_locator.json`` 由 ``backend/scripts/gen/gen_prev3_prefill_locator.py`` 扫模板生成，
每格记确切单元格、列头路径、行标签、模板该格是否公式与动作：

* ``write``  —— 目标格不是公式：写上年同底稿取数结果；取不到不写（不写 0），批注原因。
* ``verify`` —— 目标格是模板公式（多为合计）：不改写，批注上年值与来源，供核对期初数。
* ``none``   —— 模板没有对应格（原因写在定位表里），不动。

运行时的模板可能是事务所覆盖版：落格前按定位表记的列头 / 行标签 / 公式标记核对，不符就不动并记日志（不猜）。
初次预填（``wp_template_init_service.prefill_workpaper_xlsx``）与重新预填（``wp_effective.locate_cell``）共用本模块。
"""
from __future__ import annotations

import json
import logging
import re
from dataclasses import dataclass
from decimal import Decimal
from pathlib import Path
from typing import Any

logger = logging.getLogger(__name__)

BACKEND_DIR = Path(__file__).resolve().parents[2]
LOCATOR_FILE = BACKEND_DIR / "data" / "prev3_prefill_locator.json"
#: 映射表里三参 PREV 的唯一写法（192 格全是整格单个调用）
PREV3_RE = re.compile(r"^=?\s*PREV\(\s*'([^']*)'\s*,\s*'([^']*)'\s*,\s*'([^']*)'\s*\)\s*$", re.IGNORECASE)
ACTIONS = ("write", "verify", "none")
#: 列头属上年 / 期初语义的字样；比较类字样（同列出现即不是取值列）
PRIOR_TOKENS = ("上年", "上期", "期初", "年初")
COMPARISON_TOKENS = ("比较", "变动", "差异", "增长", "增减", "结构比", "比率", "比例", "占", "率")
#: 第 1、2 行是事务所名与表名；列头路径从第 3 行起
_HEADER_FIRST_ROW = 3
_HEADER_MAX_LEN = 24
_HEADER_NOISE = ("索引", "页次", "编制", "复核", "日期")
_CODE_RE = re.compile(r"^[A-Z]{1,2}\d+(-\d+)*$")
_A1_RE = re.compile(r"^([A-Z]{1,3})([0-9]{1,7})$")


@dataclass(frozen=True)
class LocatorEntry:
    wp_code: str
    sheet: str
    cell_ref: str
    prev: tuple[str, str, str]
    template: str
    cell: str | None
    header: tuple[str, ...]
    row_label: str
    cell_is_formula: bool
    action: str
    rule: str
    reason: str


# ── 模板读取（生成器与运行时核对同一套口径）────────────────────────────────────────


def norm(value: Any) -> str:
    """单元格文字去全部空白（含全角空格）；非文字返回空串。"""
    return re.sub(r"\s+", "", value) if isinstance(value, str) else ""


def is_formula_value(value: Any) -> bool:
    return isinstance(value, str) and value.lstrip().startswith("=")


def merged_values(ws: Any) -> dict[tuple[int, int], Any]:
    """合并区域内每格 → 左上角的值（列头跨列合并时每列都看得到）。"""
    out: dict[tuple[int, int], Any] = {}
    for rng in ws.merged_cells.ranges:
        value = raw_value(ws, rng.min_row, rng.min_col)
        for r in range(rng.min_row, rng.max_row + 1):
            for c in range(rng.min_col, rng.max_col + 1):
                out[(r, c)] = value
    return out


def raw_value(ws: Any, row: int, col: int) -> Any:
    """读值不建格（``ws.cell()`` 会给空格建对象，保存时多写一堆空单元格）。"""
    cell = ws._cells.get((row, col))  # noqa: SLF001 —— openpyxl 无公开的只读取格接口
    return cell.value if cell is not None else None


def _value(ws: Any, merged: dict, row: int, col: int) -> Any:
    return merged[(row, col)] if (row, col) in merged else raw_value(ws, row, col)


def header_path(ws: Any, col: int, row: int, merged: dict | None = None) -> tuple[str, ...]:
    """第 ``col`` 列在第 ``row`` 行之上（第 3 行起）的列头文字：合并单元格展开、去空白、相邻去重；
    公式、长句（> 24 字）、索引号 / 页次 / 编制复核栏、底稿编号不算列头。"""
    merged = merged_values(ws) if merged is None else merged
    parts: list[str] = []
    for r in range(_HEADER_FIRST_ROW, row):
        value = _value(ws, merged, r, col)
        if is_formula_value(value):
            continue
        text = norm(value)
        if not text or len(text) > _HEADER_MAX_LEN or any(n in text for n in _HEADER_NOISE) or _CODE_RE.match(text):
            continue
        if not parts or parts[-1] != text:
            parts.append(text)
    return tuple(parts)


def row_label(ws: Any, row: int, merged: dict | None = None) -> str:
    """行标签：A~D 列第一个非公式文字（去空白）。"""
    merged = merged_values(ws) if merged is None else merged
    for col in range(1, 5):
        value = _value(ws, merged, row, col)
        if not is_formula_value(value) and norm(value):
            return norm(value)
    return ""


def split_a1(coord: str) -> tuple[int, int]:
    """``"E13"`` → (行 13, 列 5)。"""
    from openpyxl.utils import column_index_from_string

    m = _A1_RE.match((coord or "").strip().upper())
    if m is None:
        raise ValueError(f"不是单元格坐标：{coord!r}")
    return int(m.group(2)), column_index_from_string(m.group(1))


def target_signature(ws: Any, coord: str, merged: dict | None = None) -> tuple[tuple[str, ...], str, bool]:
    """（列头路径, 行标签, 是否公式）—— 定位表记录与运行时核对用同一函数。"""
    row, col = split_a1(coord)
    merged = merged_values(ws) if merged is None else merged
    return header_path(ws, col, row, merged), row_label(ws, row, merged), is_formula_value(raw_value(ws, row, col))


def check_target(ws: Any, entry: LocatorEntry) -> str:
    """运行时模板与定位表记录一致 ⇒ 空串；否则返回不符之处（调用方不落格）。"""
    header, label, is_formula = target_signature(ws, entry.cell or "")
    if header != entry.header:
        return f"列头「{'/'.join(header)}」与定位表「{'/'.join(entry.header)}」不符"
    if label != entry.row_label:
        return f"行标签「{label}」与定位表「{entry.row_label}」不符"
    if is_formula != entry.cell_is_formula:
        return "模板该格" + ("是公式" if is_formula else "不是公式") + "，与定位表不符"
    return ""


# ── 定位表 ──────────────────────────────────────────────────────────────────────

_cache: dict[str, Any] = {"mtime": None, "entries": {}}


def _entry(wp_code: str, sheet: str, cell_ref: str, raw: dict) -> LocatorEntry:
    prev = tuple(raw.get("prev") or ())
    if len(prev) != 3 or raw.get("action") not in ACTIONS:
        raise ValueError(f"定位表条目不完整：{wp_code}|{sheet}|{cell_ref}")
    if raw["action"] != "none" and not raw.get("cell"):
        raise ValueError(f"定位表条目缺单元格：{wp_code}|{sheet}|{cell_ref}")
    return LocatorEntry(
        wp_code=wp_code, sheet=sheet, cell_ref=cell_ref, prev=prev,  # type: ignore[arg-type]
        template=raw.get("template") or "", cell=raw.get("cell") or None,
        header=tuple(raw.get("header") or ()), row_label=raw.get("row_label") or "",
        cell_is_formula=bool(raw.get("cell_is_formula")), action=raw["action"],
        rule=raw.get("rule") or "", reason=raw.get("reason") or "")


def load_locator() -> dict[tuple[str, str, str], LocatorEntry]:
    """``(wp_code, 映射 sheet, cell_ref)`` → 条目；文件缺失 / 损坏返回空（调用方据此不落格，不退回猜格）。"""
    try:
        mtime = LOCATOR_FILE.stat().st_mtime
    except OSError:
        logger.error("三参 PREV 定位表缺失：%s", LOCATOR_FILE)
        return {}
    if _cache["mtime"] == mtime:
        return _cache["entries"]
    try:
        doc = json.loads(LOCATOR_FILE.read_text(encoding="utf-8"))
        entries = {
            (wp, sheet, ref): _entry(wp, sheet, ref, raw)
            for wp, sheets in (doc.get("entries") or {}).items()
            for sheet, cells in sheets.items()
            for ref, raw in cells.items()
        }
    except (OSError, ValueError, TypeError, AttributeError) as exc:
        logger.error("三参 PREV 定位表读取失败：%s", exc)
        return {}
    _cache.update(mtime=mtime, entries=entries)
    return entries


def entries_for(wp_code: str) -> list[LocatorEntry]:
    code = (wp_code or "").strip()
    return [e for key, e in sorted(load_locator().items()) if key[0] == code]


def is_prev3_slot(wp_code: str, sheet: str, cell_ref: str) -> bool:
    """该公式位在系统预设映射里是三参 PREV（落格只认定位表）。"""
    from app.services.formula_settings.domain_stores import _mapping_records, wp_slot

    record = _mapping_records().get(wp_slot(wp_code, sheet, cell_ref))
    return record is not None and bool(PREV3_RE.match(str((record.definition or {}).get("expression") or "")))


def locate(ws: Any, wp_code: str, cell_ref: str) -> tuple[LocatorEntry | None, str]:
    """三参 PREV 公式位在本工作表上的落格 →（条目, 不落格原因）。

    条目缺失 / 动作 ``none`` / 模板与记录不符 ⇒ 条目为 ``None`` 并给原因；调用方不得退回关键词猜格。
    """
    entry = load_locator().get(((wp_code or "").strip(), ws.title.strip(), (cell_ref or "").strip()))
    if entry is None:
        return None, "定位表没有该格"
    if entry.action == "none":
        return None, entry.reason or "模板没有对应格"
    mismatch = check_target(ws, entry)
    if mismatch:
        logger.warning("三参 PREV 落格与定位表不符，不落格 wp=%s sheet=%s ref=%s：%s",
                       wp_code, ws.title, cell_ref, mismatch)
        return None, mismatch
    return entry, ""


# ── 初次预填：取值与落格 ─────────────────────────────────────────────────────────


@dataclass(frozen=True)
class PrevPrefill:
    """一格的预填内容：``value`` 为 ``None`` 时 ``reason`` 是取不到的原因（空 = 不批注）；``note`` 写来源。"""

    value: Decimal | None
    reason: str = ""
    note: str = ""
    label: str = "上年同底稿取数"


async def _source_text(db: Any, source: Any) -> str:
    if source.same_project:
        return f"本项目 {source.year} 年数据"
    from app.models.core import Project
    import sqlalchemy as sa

    try:
        async with db.begin_nested():
            name = (await db.execute(sa.select(Project.name).where(Project.id == source.project_id))).scalar_one_or_none()
    except Exception:  # noqa: BLE001 —— 只影响批注文字
        name = None
    return f"上年项目{f'「{name}」' if name else ''} {source.year} 年数据"


async def resolve_prev_prefill(db: Any, project_id: Any, wp_code: str,
                               overrides: dict[str, str] | None = None) -> dict[str, PrevPrefill]:
    """本底稿定位表各格（动作非 ``none``）的上年同底稿取数，键 ``{sheet}!{cell_ref}``。

    取值与重新预填 / 公式运行时同一来源（``PriorWpFieldResolver``）；公式位被本项目 / 全局改成别的三参 PREV 时按改后的
    参数取，改成非 PREV 公式时不在这里取（由预填按覆盖公式取值）。年度经 ``resolve_project_year``，取不到不猜今年。
    """
    entries = [e for e in entries_for(wp_code) if e.action != "none"]
    if not entries:
        return {}
    from app.services.prev_wp_field import PriorWpFieldResolver
    from app.services.sampling_registry_service import resolve_project_year

    try:
        async with db.begin_nested():
            year = await resolve_project_year(db, project_id)
    except Exception as exc:  # noqa: BLE001
        logger.warning("初次预填解析审计年度失败 project=%s: %s", project_id, exc)
        year = None
    keyed = {f"{e.sheet}!{e.cell_ref}": e for e in entries}
    if not year:
        return {k: PrevPrefill(None, "项目没有审计年度，无法确定上年") for k in keyed}
    resolver = PriorWpFieldResolver(db, project_id, int(year))
    source = await resolver.source()
    source_text = await _source_text(db, source) if source.project_id is not None else ""
    out: dict[str, PrevPrefill] = {}
    for key, entry in keyed.items():
        args = entry.prev
        if overrides and key in overrides:
            m = PREV3_RE.match(str(overrides[key] or ""))
            if m is None:
                continue
            args = m.groups()  # type: ignore[assignment]
        hit = await resolver.resolve(*args)
        if hit.value is None:
            out[key] = PrevPrefill(None, hit.reason or "上年同底稿取数取不到")
        else:
            out[key] = PrevPrefill(hit.value, note=f"{source_text}按「{hit.expression}」求值")
    return out


def slot_keys(wp_code: str) -> set[tuple[str, str]]:
    """本底稿由定位表接管落格的公式位 ``(sheet, cell_ref)``（初次预填不再按关键词猜格）。"""
    return {(e.sheet, e.cell_ref) for e in entries_for(wp_code)}


def override_prefills(mapping_data: dict, wp_code: str, overrides: dict[str, str] | None, tb_data: dict,
                      adj_data: dict | None) -> dict[str, PrevPrefill]:
    """三参 PREV 公式位被改成非 PREV 公式时，按改后的公式取试算表 / 调整数（取数口径同初次预填其余格）。"""
    if not overrides:
        return {}
    from app.services.formula_settings.wp_effective import override_value

    blocks = {(m.get("wp_code"), (m.get("sheet") or "").strip()): m for m in mapping_data.get("mappings") or []}
    out: dict[str, PrevPrefill] = {}
    for entry in entries_for(wp_code):
        key = f"{entry.sheet}!{entry.cell_ref}"
        formula = overrides.get(key)
        if formula is None or PREV3_RE.match(str(formula)):
            continue
        block = blocks.get((wp_code, entry.sheet)) or {}
        preset = next((c.get("formula", "") for c in block.get("cells") or [] if c.get("cell_ref") == entry.cell_ref), "")
        value = override_value(formula, preset, tb_data, adj_data, block.get("account_codes") or [])
        out[key] = PrevPrefill(None if value is None else Decimal(str(value)), note=f"按「{formula}」取数",
                               label="改写后的公式取数")
    return out


def _annotate(cell: Any, text: str) -> None:
    """追加批注（模板原有的编制说明保留在前）。"""
    from openpyxl.comments import Comment

    old = cell.comment
    cell.comment = Comment(f"{old.text}\n\n{text}" if old is not None and old.text else text, "系统")


def apply_prefills(wb: Any, wp_code: str, prefills: dict[str, PrevPrefill],
                   skip_cells: set[tuple[str, str]] | frozenset = frozenset()) -> int:
    """按定位表把 ``prefills``（键 ``{sheet}!{cell_ref}``）落到工作簿，返回写值格数。

    ``write`` 写值 + 浅蓝底 + 批注来源；``verify`` 只批注（模板公式不改写）；取不到只批注原因、不写 0；模板与定位表
    不符不动；同一目标格只落一次；``skip_cells``（用户自定义公式所在格）不动。
    """
    from openpyxl.styles import PatternFill

    sheets = {ws.title.strip(): ws for ws in wb.worksheets}
    done: set[tuple[str, str]] = set()
    written = 0
    for entry in entries_for(wp_code):
        prefill = prefills.get(f"{entry.sheet}!{entry.cell_ref}")
        ws = sheets.get(entry.sheet)
        if prefill is None or ws is None or entry.action == "none":
            continue
        located, _why = locate(ws, wp_code, entry.cell_ref)
        target = (entry.sheet, located.cell if located else "")
        if located is None or target in done or target in skip_cells:
            continue
        done.add(target)
        cell = ws[located.cell]
        if prefill.value is None:
            if prefill.reason:
                _annotate(cell, f"{prefill.label}取不到：{prefill.reason}")
            continue
        amount = f"{prefill.value:,.2f}"
        if located.action == "write":
            cell.value = float(prefill.value)
            cell.fill = PatternFill(start_color="E8F4FD", end_color="E8F4FD", fill_type="solid")
            _annotate(cell, f"系统预填充：{prefill.label} {amount}（{prefill.note}）")
            written += 1
        else:
            _annotate(cell, f"{prefill.label}核对：{amount}（{prefill.note}）。本格为模板公式，未改写。")
    return written

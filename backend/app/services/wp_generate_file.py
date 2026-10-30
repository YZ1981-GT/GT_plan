"""新建底稿的文件落地（按编码生成 / 准则转换生成共用）。

spec: .kiro/specs/_archive/71-formula-fetch-and-wp-generation-defect-closure/（R5）

与链路编排（``chain_orchestrator``）同一解析：``init_workpaper_from_template`` 按底稿模板目录索引找主文件
（``find_template_file_any``，支持「E1-1至E1-11 …」这类区间命名与 xlsm / docx），多页底稿合并其余文件的
sheet，并按项目场景裁剪（normal 不带 IPO / 上市 / 重组 / 舞弊应对文件）。文件落在
``wp_storage/{project_id}/{wp_id}.*`` —— 与在线编辑首开初始化、链路编排是同一位置。

🔴 此前按编码生成与准则转换生成各自按 ``gt_template_library.json`` 精确匹配：模板集 1025 个编码里
722 个不在该库（E1 / D2 等多页底稿全在其中），统统落成只有三行字的空白册，且名称叫「底稿{code}」。
找不到模板时本模块仍建空白册（同一存储位置），由调用方单列回报，不冒充模板。
"""
from __future__ import annotations

import logging
import re
from dataclasses import dataclass
from pathlib import Path
from uuid import UUID

import sqlalchemy as sa
from sqlalchemy.ext.asyncio import AsyncSession

from app.services import wp_template_init_service as _init

logger = logging.getLogger(__name__)

#: Excel 页签名不允许的字符（openpyxl 遇到会抛错）
_SHEET_TITLE_BAD = re.compile(r"[\[\]:*?/\\]")

#: 找不到模板时空白册里的说明文字与所在行
BLANK_NOTICE = "未找到该编码的底稿模板，这是一份空白底稿"
BLANK_NOTICE_ROW = 8


@dataclass(frozen=True)
class GeneratedWorkpaperFile:
    """``path``：底稿文件；``from_template``：是否来自真实模板（False = 找不到模板建的空白册）。"""

    path: Path
    from_template: bool


async def load_generation_scenario(db: AsyncSession, project_id: UUID) -> tuple[str, bool]:
    """项目场景与是否有外币业务（与链路编排同口径；取不到按 normal / 无外币）。

    场景取 ``project_classification`` 单一解释（R4）：迁移前残留在分类列里的 IPO / refinancing
    在场景仍为 normal 时按场景处理，生成文件的场景裁剪与其他模块一致。
    """
    from app.models.core import Project
    from app.services.project_classification import resolve

    row = (await db.execute(
        sa.select(Project.business_category, Project.scenario, Project.has_foreign_currency)
        .where(Project.id == project_id)
    )).first()
    if row is None:
        return "normal", False
    return resolve(row[0], row[1]).scenario, bool(row[2])


def materialize_workpaper_file(
    project_id: UUID,
    wp_id: UUID,
    wp_code: str,
    wp_name: str,
    *,
    scenario: str = "normal",
    has_foreign_currency: bool = False,
    target_base_path: Path | None = None,
) -> GeneratedWorkpaperFile:
    """按模板生成底稿文件；找不到模板时在同一存储位置建空白册并标 ``from_template=False``。"""
    path = _init.init_workpaper_from_template(
        project_id=project_id,
        wp_id=wp_id,
        wp_code=wp_code,
        scenario=scenario,
        has_foreign_currency=has_foreign_currency,
        target_base_path=target_base_path,
    )
    if path is not None:
        return GeneratedWorkpaperFile(path=Path(path), from_template=True)

    import openpyxl

    target = target_base_path or _init.get_workpaper_storage_path(project_id, wp_id)
    target.parent.mkdir(parents=True, exist_ok=True)
    wb = openpyxl.Workbook()
    try:
        ws = wb.active
        ws.title = (_SHEET_TITLE_BAD.sub("_", wp_code or "") or "Sheet1")[:31]
        ws["A1"] = f"底稿编号: {wp_code}"
        ws["A2"] = f"底稿名称: {wp_name}"
        # 第 8 行：表头填充（wp_header_service 标准布局占第 1~5 行）不会覆盖它
        ws[f"A{BLANK_NOTICE_ROW}"] = BLANK_NOTICE
        wb.save(str(target))
    finally:
        wb.close()
    logger.warning("生成底稿：未找到模板，已建空白底稿 wp_code=%s wp_id=%s", wp_code, wp_id)
    return GeneratedWorkpaperFile(path=target, from_template=False)


def discard_generated_file(generated: GeneratedWorkpaperFile | None) -> None:
    """生成流程在落文件之后失败（底稿行已回滚）时删掉孤儿文件；删不掉只记日志。"""
    if generated is None:
        return
    try:
        generated.path.unlink(missing_ok=True)
    except OSError as exc:
        logger.warning("生成底稿：清理孤儿文件失败 %s: %s", generated.path, exc)


async def create_template_workpaper(
    db: AsyncSession,
    receipt,
    *,
    project_id: UUID,
    wp_code: str,
    wp_name: str,
    audit_cycle: str | None,
    scenario: str = "normal",
    has_foreign_currency: bool = False,
):
    """建一份模板底稿（索引行 + 底稿行 + 文件）并记进回执（``GenerationReceipt``）；链路编排 / IPO / B60 共用。

    单码保存点：建行或落文件失败只回滚这一码、删掉已落的文件、记入 ``receipt.failures``，不计入新建
    （spec procedure-chain-hardening-and-java-contracts R11.3 / R11.4；原先复制模板在保存点外、失败照样计入新建，
    链路编排单码失败还整会话 rollback）。找不到模板建空白册；请求内不填表头（计入 ``header_unfilled``）。
    返回新建的 ``WorkingPaper``，失败返回 None。调用方先经 ``plan_generation`` 判定保留范围。
    """
    from app.models.workpaper_models import WorkingPaper, WpIndex, WpSourceType, WpStatus

    generated: GeneratedWorkpaperFile | None = None
    try:
        async with db.begin_nested():
            wp_index = WpIndex(project_id=project_id, wp_code=wp_code, wp_name=wp_name,
                               audit_cycle=audit_cycle, status=WpStatus.not_started)
            db.add(wp_index)
            await db.flush()
            wp = WorkingPaper(wp_index_id=wp_index.id, project_id=project_id,
                              source_type=WpSourceType.template, file_path="", parsed_data={})
            db.add(wp)
            await db.flush()  # 取底稿 id：文件按 wp_storage/{项目}/{底稿 id} 存放
            generated = materialize_workpaper_file(project_id, wp.id, wp_code, wp_name, scenario=scenario,
                                                   has_foreign_currency=has_foreign_currency)
            wp.file_path = str(generated.path)
            await db.flush()
    except Exception as exc:  # noqa: BLE001 — 单码失败不影响其余，如实进回执
        discard_generated_file(generated)
        receipt.add_failure(wp_code, exc)
        logger.warning("生成底稿失败 wp_code=%s: %s", wp_code, exc, exc_info=True)
        return None
    except BaseException:
        discard_generated_file(generated)
        raise
    receipt.add_created(wp_code, generated, header_filled=False)
    return wp


@dataclass(frozen=True)
class WorkpaperFileJob:
    """批量落文件的一项：``library_file`` = 项目选定模板；``target_base_path`` = 可选 staging 基础路径。"""

    wp_id: UUID
    wp_code: str
    wp_name: str
    library_file: str | None = None
    target_base_path: Path | None = None


def _copy_library_file(project_id: UUID, job: WorkpaperFileJob) -> GeneratedWorkpaperFile | None:
    """项目选定的模板库文件存在 ⇒ 复制一份到本底稿存储位置（扩展名照源文件）；不存在返回 None。"""
    if not job.library_file:
        return None
    from app.services.wp_export.wp_file_resolver import resolve_wp_file

    source = resolve_wp_file(job.library_file, wp_code=None, allow_template_fallback=False).path
    if source is None:
        logger.warning("生成底稿：项目选定的模板文件不存在，改按模板目录 wp_code=%s file=%s",
                       job.wp_code, job.library_file)
        return None
    import shutil

    target_base = job.target_base_path or _init.get_workpaper_storage_path(project_id, job.wp_id)
    target_base.parent.mkdir(parents=True, exist_ok=True)
    target = target_base.with_suffix(source.suffix)
    shutil.copy2(source, target)
    return GeneratedWorkpaperFile(path=target, from_template=True)


def materialize_workpaper_files(
    project_id: UUID,
    jobs: list[WorkpaperFileJob],
    *,
    scenario: str = "normal",
    has_foreign_currency: bool = False,
) -> list[GeneratedWorkpaperFile]:
    """整批落文件（与 ``materialize_workpaper_file`` 同一解析，顺序与 ``jobs`` 一致）。

    全有或全无：任何一份失败 ⇒ 删掉本批已落的文件再抛出（调用方的底稿行随事务回滚）。不读写数据库，
    调用方可放到线程里跑（重药安徽 737 个编码实测约 31 秒）。
    """
    done: list[GeneratedWorkpaperFile] = []
    try:
        for job in jobs:
            generated = _copy_library_file(project_id, job) or materialize_workpaper_file(
                project_id, job.wp_id, job.wp_code, job.wp_name,
                scenario=scenario, has_foreign_currency=has_foreign_currency,
                target_base_path=job.target_base_path,
            )
            done.append(generated)
    except BaseException:
        for generated in done:
            discard_generated_file(generated)
        raise
    return done


def detach_from_shared_template(wp, template_path: Path) -> Path:
    """在线保存时底稿 ``file_path`` 指向共享模板目录 ⇒ 先复制到本底稿存储位置并改写 ``file_path``（copy-on-write）。

    🔴 生成底稿曾直接引用模板目录里的文件（如 H1 ⇒ ``wp_templates/H/H1 固定资产.xlsx``）。模板目录在 legacy 路径
    边界之内（backend 根下），就地保存会把一个项目的编辑写进所有项目共用的模板（spec procedure-chain-followup-closure
    R3.3）。模板目录永不写：复制一份（目标已存在就沿用），之后读写都走副本。只改 ``wp.file_path``，由调用方提交。
    """
    import shutil

    target = _init.get_workpaper_storage_path(wp.project_id, wp.id).with_suffix(template_path.suffix)
    if not target.exists():
        shutil.copy2(template_path, target)
    logger.warning("在线保存：底稿文件指向共享模板，已复制到本底稿存储 wp=%s %s → %s", wp.id, template_path, target)
    wp.file_path = str(target)
    return target


def needs_own_file(raw_file_path: str | None) -> bool:
    """``file_path`` 还没指向一份本底稿自己的文件：为空、指向共享模板目录、或指向的文件不存在 / 越界。"""
    from app.services.workpaper_sync.canonical_paths import TEMPLATE_ROOT, is_within_root
    from app.services.wp_export.wp_file_resolver import resolve_wp_file

    resolved = resolve_wp_file(raw_file_path, wp_code=None, allow_template_fallback=False).path
    return resolved is None or is_within_root(TEMPLATE_ROOT, resolved)


async def bind_lazy_workpaper_file(db: AsyncSession, wp_id: UUID, path: Path) -> bool:
    """打开底稿时懒加载落了本底稿文件 ⇒ 回写 ``file_path``（spec procedure-chain-followup-closure R3.3）。

    只在原值还不是本底稿自己的文件时改（见 ``needs_own_file``）；原值指向一份存在的本底稿文件（如 legacy
    ``storage/projects/...``）不动 —— 改了会让那份带编辑内容的文件失去引用。条件更新：并发的另一次打开先写了，
    本次 0 行。返回是否改写；只执行语句不提交，由调用方 commit。
    """
    from app.models.workpaper_models import WorkingPaper

    row = (await db.execute(
        sa.select(WorkingPaper.file_path).where(WorkingPaper.id == wp_id)
    )).first()
    if row is None or not needs_own_file(row[0]):
        return False
    result = await db.execute(
        sa.update(WorkingPaper)
        .where(WorkingPaper.id == wp_id, WorkingPaper.file_path.is_not_distinct_from(row[0]))
        .values(file_path=str(path))
        .execution_options(synchronize_session=False)
    )
    if result.rowcount:
        logger.info("底稿懒加载落了文件，已回写 file_path wp=%s %r → %s", wp_id, row[0], path)
    return bool(result.rowcount)

"""BulkExportService — 项目级底稿批量导出编排（Req 1/3/6）

编排流程：
  build_manifest → for each exportable entry → export_tab → ZipAssembler.write
  → 写 manifest.json（含 sha256）+ README.txt → finalize

fail-soft：单 Tab 导出失败 → log warning, skip that tab, continue with rest。
mode=data + only_with_data：空表（无数据行）标 skipped("no_data")。
service 只 flush 不 commit。

Requirements: 1.1, 1.2, 1.7, 3.1, 3.2, 3.3, 6.3
"""
from __future__ import annotations

import hashlib
import io
import json
import logging
from typing import Any, Literal, Protocol
from uuid import UUID

from sqlalchemy.ext.asyncio import AsyncSession

from app.services.bulk_tab.manifest_builder import (
    BulkManifest,
    ManifestFileEntry,
    ManifestSkippedEntry,
    build_manifest,
)
from app.services.bulk_tab.single_tab_adapter import export_tab
from app.services.bulk_tab.zip_handler import ZipAssembler

logger = logging.getLogger(__name__)


# ---------------------------------------------------------------------------
# Progress callback protocol
# ---------------------------------------------------------------------------


class ProgressCallback(Protocol):
    """SSE 进度回调协议。"""

    def tick(self, message: str = "") -> None: ...


# ---------------------------------------------------------------------------
# Helper: SHA-256
# ---------------------------------------------------------------------------


def _sha256(data: bytes) -> str:
    """计算 bytes 的 SHA-256 hex digest。"""
    return hashlib.sha256(data).hexdigest()


# ---------------------------------------------------------------------------
# Helper: 检测空 xlsx（mode=data + only_with_data）
# ---------------------------------------------------------------------------


def _is_empty_xlsx(xlsx_bytes: bytes) -> bool:
    """检测 xlsx 文件是否为空数据文件（仅有表头、无数据行）。

    策略：用 openpyxl 只读模式打开，检查第一个 sheet 的数据行数。
    若仅有 0~1 行（可能是表头）则视为空。

    对于无法解析的文件（非 xlsx），保守返回 False（不跳过）。
    """
    try:
        import openpyxl

        wb = openpyxl.load_workbook(
            io.BytesIO(xlsx_bytes), read_only=True, data_only=True
        )
        ws = wb.active
        if ws is None:
            wb.close()
            return True

        # 统计非空行数（跳过表头行）
        row_count = 0
        for i, row in enumerate(ws.iter_rows(values_only=True)):
            if i == 0:
                # 第一行通常是表头
                continue
            # 判断该行是否有任何非空值
            if any(cell is not None and str(cell).strip() != "" for cell in row):
                row_count += 1
                if row_count > 0:
                    # 有至少 1 行数据即非空
                    wb.close()
                    return False

        wb.close()
        return row_count == 0
    except Exception:
        # 解析失败保守不跳过
        logger.debug("_is_empty_xlsx: 无法解析 xlsx，保守返回 False")
        return False


# ---------------------------------------------------------------------------
# Helper: README.txt 渲染
# ---------------------------------------------------------------------------


def _render_readme(manifest: BulkManifest) -> str:
    """生成 README.txt 内容（Req 1.7）。

    包含：
    - 导出摘要（各循环状态统计）
    - ZIP 使用说明
    - 导出信息（项目/模式/循环/导出时间）
    - 逐 Tab 文件列表（含编制提示链接引导）
    - 附加文件说明
    """
    mode_label = "模板" if manifest.mode == "template" else "数据"
    lines: list[str] = [
        "=" * 60,
        "审计底稿批量导出",
        "=" * 60,
        "",
        f"导出时间: {manifest.exported_at}",
        f"导出人: {manifest.exported_by}",
        f"审计年度: {manifest.audit_year}",
        f"平台版本: {manifest.platform_version}",
        "",
        "─" * 40,
        "导出摘要",
        "─" * 40,
        "",
    ]

    # Per-cycle summary
    from collections import Counter
    cycle_counts: Counter[str] = Counter()
    for f in manifest.files:
        # Extract cycle from zip_path first segment
        parts = f.zip_path.split("/")
        cycle_counts[parts[0] if parts else "?"] += 1

    lines.append(f"已导出底稿: {len(manifest.files)} 张")
    lines.append(f"跳过: {len(manifest.skipped)} 张")
    lines.append("")
    lines.append("各循环状态:")
    for cycle in sorted(cycle_counts.keys()):
        count = cycle_counts[cycle]
        skip_count = sum(1 for s in manifest.skipped if s.parent_wp_code.startswith(cycle))
        lines.append(f"  {cycle}: 导出 {count} 张" + (f", 跳过 {skip_count}" if skip_count else ""))

    if manifest.skipped:
        lines.append("")
        lines.append("跳过原因:")
        for s in manifest.skipped[:20]:  # limit display
            lines.append(f"  {s.sheet_code}: {s.skip_reason}")
        if len(manifest.skipped) > 20:
            lines.append(f"  ...（共 {len(manifest.skipped)} 项，仅显示前 20 项）")

    lines.append("")
    lines.append("─" * 40)
    lines.append("附加文件")
    lines.append("─" * 40)
    lines.append("  _报表/财务报表.xlsx — 四张财务报表（审定数）")
    lines.append("  _报表/财务报表_未审数.xlsx — 四张财务报表（未审数，如支持）")
    lines.append("  _附注/财务报表附注.docx — 附注正文")
    lines.append("  _试算表/试算平衡表.xlsx — 科目余额试算表")
    lines.append("  _底稿目录.xlsx — 底稿清单索引")
    lines.append("  （以上文件按项目数据就绪状态自动纳入，缺失即该项暂无数据）")
    lines.append("")
    lines.append("─" * 40)
    lines.append("使用说明")
    lines.append("─" * 40)
    lines.append("")

    if manifest.mode == "template":
        lines.extend([
            "1. 解压本 ZIP 到本地文件夹。",
            "2. 按需打开各 Excel 文件，参照编制提示填写数据。",
            "3. 填写完毕后，将整个文件夹重新压缩为 ZIP（保持目录结构不变）。",
            "4. 在系统中选择「导入全部数据」上传 ZIP 即可批量导入。",
            "",
            "注意事项：",
            "- 请勿修改文件名或目录结构，否则导入时将无法识别。",
            "- 请勿修改表头行（第一行），仅在数据行填写。",
            "- 支持部分导入：未修改的文件可删除，系统会跳过缺失文件。",
        ])
    else:
        lines.extend([
            "1. 各 xlsx 文件可直接在 Excel 中打开编辑。",
            "2. 编辑后可通过「批量导入」功能回写平台。",
            "3. manifest.json 包含完整导出清单和文件哈希。",
        ])

    lines.extend([
        "",
        "─" * 40,
        "导出信息",
        "─" * 40,
        "",
        f"- 项目ID: {manifest.project_id}",
        f"- 审计年度: {manifest.audit_year}",
        f"- 导出模式: {mode_label}",
        f"- 审计循环: {', '.join(manifest.cycles)}",
        f"- 导出时间: {manifest.exported_at}",
        f"- 导出人: {manifest.exported_by}",
        f"- 平台版本: {manifest.platform_version}",
        "",
        "─" * 40,
        "文件清单",
        "─" * 40,
        "",
    ])

    for f in manifest.files:
        if f.sha256:
            lines.append(f"- [{f.sheet_code}] {f.sheet_name}")
            lines.append(f"  路径: {f.zip_path}")
            lines.append("")

    lines.append("")
    lines.append("---")
    lines.append("本文件由系统自动生成，请勿手动修改。")
    lines.append("")

    return "\n".join(lines)


# ---------------------------------------------------------------------------
# Main export function
# ---------------------------------------------------------------------------


async def export(
    db: AsyncSession,
    project_id: UUID,
    cycles: list[str] | None,
    mode: Literal["template", "data"],
    *,
    only_with_data: bool = False,
    incremental: bool = False,
    password: str | None = None,
    exported_by: str = "",
    platform_version: str = "",
    audit_year: int = 0,
    progress: Any | None = None,
    visible_filter: Any | None = None,
) -> io.BytesIO:
    """批量导出底稿为 ZIP。

    编排流程：
    1. build_manifest 获取 BulkManifest
    2. 遍历 manifest.exportable() 逐 Tab 调 export_tab
    3. mode=data + only_with_data 时检测空表跳过
    4. 计算 sha256 更新 entry
    5. 写 manifest.json + README.txt
    6. finalize 返回 BytesIO

    fail-soft: 单 Tab 失败 → log warning, skip, continue。
    service 只 flush 不 commit。

    Args:
        db: AsyncSession（只 flush 不 commit）。
        project_id: 项目 UUID。
        cycles: 审计循环列表（如 ["D"]）。None 导出全部。
        mode: "template" 或 "data"。
        only_with_data: mode=data 时是否仅导出有数据的 Tab（Req 3.3）。
        exported_by: 导出人用户名。
        platform_version: 平台版本号。
        audit_year: 审计年度。
        progress: SSE 进度回调（可选）。

    Returns:
        BytesIO 包含完整 ZIP。

    Raises:
        AcnrCatalogUnavailableError: ACNR catalog 不可用。
    """
    # ─── Step 1: Build manifest ───────────────────────────────────────
    manifest = await build_manifest(
        db=db,
        project_id=project_id,
        cycles=cycles,
        mode=mode,
        exported_by=exported_by,
        platform_version=platform_version,
        audit_year=audit_year,
    )

    # ─── Step 2: Create ZipAssembler ──────────────────────────────────
    zip_assembler = ZipAssembler()

    # ─── Step 3: Export each tab ──────────────────────────────────────
    exportable_entries = manifest.exportable()

    # ─── Incremental: load previous hashes ────────────────────────────
    prev_hashes: dict[str, str] = {}
    if incremental:
        try:
            from app.models.core import Project
            proj = await db.get(Project, project_id)
            if proj and proj.wizard_state:
                prev_manifest = proj.wizard_state.get('_last_bulk_export_manifest', {})
                prev_hashes = prev_manifest.get('file_hashes', {})
        except Exception:
            pass  # fall back to full export

    # Task 10 · 可见集过滤（Req 8.6/8.14/9）：manifest 只由 gate 可见集构建。不可见底稿在 export_tab
    # （读正文）之前被剔除，其 sha256 保持空 → Step 4 从 manifest.files 移除，不泄露存在性。
    if visible_filter is not None:
        _visible: list[ManifestFileEntry] = []
        for _e in exportable_entries:
            try:
                _ok = await visible_filter(_e.wp_id, _e.sheet_code)
            except Exception:  # noqa: BLE001 — 过滤异常 fail-closed（不导出该条目）
                _ok = False
            if _ok:
                _visible.append(_e)
        exportable_entries = _visible

    # ─── Progress total tracking (Improvement #12) ────────────────────
    total_items = len(exportable_entries) + 3  # +3 for reports/notes/tb
    if progress and hasattr(progress, 'set_total'):
        progress.set_total(total_items)

    exported_count = 0
    skipped_count = 0

    for entry in exportable_entries:
        try:
            # 调用 SingleTabIeAdapter.export_tab
            xlsx_bytes = await export_tab(
                db=db,
                wp_id=entry.wp_id,
                api_prefix=entry.api_prefix,
                sheet_code=entry.sheet_code,
                mode=mode,
            )
        except KeyError:
            # api_prefix 未注册 → skip_reason=no_adapter（fail-soft）
            logger.warning(
                "bulk_export: skip %s — no_adapter (api_prefix=%s)",
                entry.sheet_code,
                entry.api_prefix,
            )
            manifest.skipped.append(
                ManifestSkippedEntry(
                    addr_id=entry.addr_id,
                    sheet_code=entry.sheet_code,
                    sheet_name=entry.sheet_name,
                    parent_wp_code=entry.parent_wp_code,
                    skip_reason="no_adapter",
                )
            )
            # 从 files 中移除此 entry（通过清空 sha256 标记，后续不写入）
            entry.sha256 = ""
            skipped_count += 1
            if progress:
                progress.tick(f"跳过 {entry.sheet_code}（无适配器）")
            continue
        except Exception as exc:
            # 其他异常 → fail-soft，记日志跳过
            logger.warning(
                "bulk_export: skip %s — export_tab failed: %s",
                entry.sheet_code,
                exc,
                exc_info=True,
            )
            manifest.skipped.append(
                ManifestSkippedEntry(
                    addr_id=entry.addr_id,
                    sheet_code=entry.sheet_code,
                    sheet_name=entry.sheet_name,
                    parent_wp_code=entry.parent_wp_code,
                    skip_reason="export_failed",
                )
            )
            entry.sha256 = ""
            skipped_count += 1
            if progress:
                progress.tick(f"跳过 {entry.sheet_code}（导出失败）")
            continue

        # ─── mode=data + only_with_data → 检测空表 ────────────────────
        if mode == "data" and only_with_data and _is_empty_xlsx(xlsx_bytes):
            logger.info(
                "bulk_export: skip %s — no_data (empty xlsx)",
                entry.sheet_code,
            )
            manifest.skipped.append(
                ManifestSkippedEntry(
                    addr_id=entry.addr_id,
                    sheet_code=entry.sheet_code,
                    sheet_name=entry.sheet_name,
                    parent_wp_code=entry.parent_wp_code,
                    skip_reason="no_data",
                )
            )
            entry.sha256 = ""
            skipped_count += 1
            if progress:
                progress.tick(f"跳过 {entry.sheet_code}（无数据）")
            continue

        # ─── 计算 sha256 并写入 ZIP ───────────────────────────────────
        file_hash = _sha256(xlsx_bytes)

        # ─── Incremental: skip if hash unchanged ──────────────────────
        if incremental and prev_hashes.get(entry.sheet_code) == file_hash:
            entry.sha256 = ""  # mark as skipped
            manifest.skipped.append(
                ManifestSkippedEntry(
                    addr_id=entry.addr_id,
                    sheet_code=entry.sheet_code,
                    sheet_name=entry.sheet_name,
                    parent_wp_code=entry.parent_wp_code,
                    skip_reason="unchanged",
                )
            )
            skipped_count += 1
            if progress:
                progress.tick(f"跳过 {entry.sheet_code}（未变更）")
            continue

        entry.sha256 = file_hash

        zip_assembler.write(entry.zip_path, xlsx_bytes)
        del xlsx_bytes  # Release memory immediately after writing to ZIP buffer
        exported_count += 1

        if progress:
            progress.tick(f"已导出 {entry.sheet_code} ({exported_count})")

    # ─── Step 4: 从 manifest.files 中移除未成功导出的 entries ─────────
    # 保留 sha256 非空的 entries 在 files 中（已成功导出）
    manifest.files = [f for f in manifest.files if f.sha256]

    # ─── Step 3b: Export reports + notes + trial balance ─────────────────
    try:
        from app.services.report_excel_exporter import ReportExcelExporter
        report_exporter = ReportExcelExporter(db)
        # Export audited version
        report_bytes_io = await report_exporter.export(
            project_id=project_id, year=audit_year, flatten_formulas=True
        )
        report_bytes = report_bytes_io.getvalue()
        zip_assembler.write("_报表/财务报表.xlsx", report_bytes)
        exported_count += 1
        # Export unadjusted version if mode parameter is supported
        import inspect
        _export_sig = inspect.signature(report_exporter.export)
        if 'mode' in _export_sig.parameters:
            try:
                report_unadj_io = await report_exporter.export(
                    project_id=project_id, year=audit_year, flatten_formulas=True, mode="unadjusted"
                )
                zip_assembler.write("_报表/财务报表_未审数.xlsx", report_unadj_io.getvalue())
            except Exception as unadj_exc:
                logger.debug("bulk_export: unadjusted report failed: %s", unadj_exc)
    except Exception as exc:
        logger.warning("bulk_export: skip reports — %s", exc)

    try:
        from app.services.note_word_exporter import NoteWordExporter
        note_exporter = NoteWordExporter(db)
        # Determine template_type from project
        from app.models.core import Project
        project = await db.get(Project, project_id)
        template_type = getattr(project, 'template_type', 'soe') or 'soe'
        note_bytes_io = await note_exporter.export(
            project_id=project_id, year=audit_year, template_type=template_type
        )
        note_bytes = note_bytes_io.getvalue()
        zip_assembler.write("_附注/财务报表附注.docx", note_bytes)
        exported_count += 1
    except Exception as exc:
        logger.warning("bulk_export: skip notes — %s", exc)

    try:
        # Trial balance export as xlsx
        from app.services.trial_balance_service import TrialBalanceService
        import io as _io
        from openpyxl import Workbook as _Wb
        tb_svc = TrialBalanceService(db)
        tb_rows = await tb_svc.get_trial_balance(project_id, audit_year)
        if tb_rows:
            wb = _Wb()
            ws = wb.active
            ws.title = "试算平衡表"
            ws.append(["科目编码", "科目名称", "期初余额", "未审数", "审计调整(AJE)", "重分类调整(RJE)", "审定数"])
            for r in tb_rows:
                ws.append([
                    getattr(r, 'standard_account_code', ''),
                    getattr(r, 'account_name', ''),
                    round(float(getattr(r, 'opening_balance', 0) or 0), 2),
                    round(float(getattr(r, 'unadjusted_amount', 0) or 0), 2),
                    round(float(getattr(r, 'aje_adjustment', 0) or 0), 2),
                    round(float(getattr(r, 'rje_adjustment', 0) or 0), 2),
                    round(float(getattr(r, 'audited_amount', 0) or 0), 2),
                ])
            buf = _io.BytesIO()
            wb.save(buf)
            zip_assembler.write("_试算表/试算平衡表.xlsx", buf.getvalue())
            exported_count += 1
    except Exception as exc:
        logger.warning("bulk_export: skip trial_balance — %s", exc)

    # ─── Step 3c: Write index sheet (_底稿目录.xlsx) ────────────────────────
    try:
        import io as _io2
        from openpyxl import Workbook as _Wb2
        from openpyxl.styles import Font as _Font2
        from openpyxl.utils import get_column_letter as _gcl2

        wb_idx = _Wb2()
        ws_idx = wb_idx.active
        ws_idx.title = "底稿目录"
        idx_headers = ["序号", "循环", "底稿编码", "科目名称", "状态", "文件路径"]
        ws_idx.append(idx_headers)
        # Bold headers
        for col_i in range(1, len(idx_headers) + 1):
            ws_idx.cell(row=1, column=col_i).font = _Font2(bold=True)

        for seq, f in enumerate(manifest.files, start=1):
            ws_idx.append([
                seq,
                f.zip_path.split("/")[0] if "/" in f.zip_path else "",
                f.sheet_code,
                f.sheet_name,
                "已导出" if f.sha256 else "跳过",
                f.zip_path,
            ])

        # Also list skipped entries
        seq = len(manifest.files)
        for s in manifest.skipped:
            seq += 1
            ws_idx.append([
                seq,
                s.parent_wp_code,
                s.sheet_code,
                s.sheet_name,
                f"跳过({s.skip_reason})",
                "",
            ])

        # Auto-width
        for col_i in range(1, len(idx_headers) + 1):
            max_len = max(
                (len(str(ws_idx.cell(row=r, column=col_i).value or "")) for r in range(1, ws_idx.max_row + 1)),
                default=10,
            )
            ws_idx.column_dimensions[_gcl2(col_i)].width = min(max_len + 2, 50)

        buf_idx = _io2.BytesIO()
        wb_idx.save(buf_idx)
        zip_assembler.write("_底稿目录.xlsx", buf_idx.getvalue())
    except Exception as exc:
        logger.warning("bulk_export: skip index sheet — %s", exc)

    # ─── Step 5: Write manifest.json ──────────────────────────────────
    manifest_json = json.dumps(
        manifest.to_dict(),
        ensure_ascii=False,
        indent=2,
    ).encode("utf-8")
    zip_assembler.write("manifest.json", manifest_json)

    # ─── Step 6: Write README.txt ─────────────────────────────────────
    readme_text = _render_readme(manifest).encode("utf-8")
    zip_assembler.write("README.txt", readme_text)

    # ─── Step 7: Finalize ─────────────────────────────────────────────
    logger.info(
        "bulk_export: project=%s mode=%s cycles=%s exported=%d skipped=%d",
        project_id,
        mode,
        cycles,
        exported_count,
        skipped_count,
    )

    # ─── Step 8: Persist manifest for incremental export ─────────────
    try:
        from app.models.core import Project
        from sqlalchemy.orm.attributes import flag_modified
        proj = await db.get(Project, project_id)
        if proj:
            ws = proj.wizard_state or {}
            ws['_last_bulk_export_manifest'] = {
                'exported_at': manifest.exported_at,
                'mode': manifest.mode,
                'file_hashes': {f.sheet_code: f.sha256 for f in manifest.files if f.sha256},
            }
            proj.wizard_state = {**ws}  # trigger ORM dirty
            flag_modified(proj, 'wizard_state')
            await db.flush()
    except Exception as exc:
        logger.warning("bulk_export: failed to persist manifest for incremental — %s", exc)

    # ─── Step 9: Optional password protection ─────────────────────────
    zip_result = zip_assembler.finalize()
    if password:
        try:
            import pyzipper
            import zipfile
            plain_bytes = zip_result.getvalue()
            encrypted_buf = io.BytesIO()
            with pyzipper.AESZipFile(encrypted_buf, 'w', compression=pyzipper.ZIP_DEFLATED, encryption=pyzipper.WZ_AES) as zf:
                zf.setpassword(password.encode('utf-8'))
                with zipfile.ZipFile(io.BytesIO(plain_bytes)) as src:
                    for name in src.namelist():
                        zf.writestr(name, src.read(name))
            encrypted_buf.seek(0)
            return encrypted_buf
        except ImportError:
            logger.warning("bulk_export: pyzipper not available, exporting without password")
        except Exception as exc:
            logger.warning("bulk_export: password protection failed — %s", exc)

    return zip_result

"""公式模块导入导出服务（Req 23）。

提供公式模块「导出模板 / 导出数据 / 导入数据」三能力的核心逻辑，被
``routers/formula_import_export.py`` 的三端点消费：

- ``build_template_workbook()``：生成含**编报说明首区块**（单一源
  ``reporting_instructions``）+ 空白/示例公式行的模板工作簿（Req 23.2）。
- ``build_data_workbook(entries)``：把当前页面/模块已有公式（预设库条目）导出为
  工作簿，含目标单元/表达式/公式类型/引用/最近计算时间（Req 23.3）。
- ``parse_import_rows(content)``：解析上传 xlsx → 每条公式引用经 ACNR
  ``full_resolve`` 校验，悬空引用报告并跳过（不静默入库错误公式，Req 23.4）。

引用一律经 ``resolve_ref``（ACNR ``full_resolve``，fail-open）校验，禁裸坐标串
（Req 11）。基础设施故障（fail-open）不视为悬空，不跳过（Req 11.3）。

Requirements: 23.1, 23.2, 23.3, 23.4, 23.5, 23.6
"""

from __future__ import annotations

import io
import json
import logging
from dataclasses import dataclass, field
from typing import Any

from openpyxl import Workbook, load_workbook
from openpyxl.styles import Alignment, Font
from openpyxl.utils import get_column_letter
from sqlalchemy.ext.asyncio import AsyncSession

from app.services.formula_management.engine import resolve_ref
from app.services.formula_management.preset_library import (
    VALID_FORMULA_TYPES,
    PresetEntry,
    _formula_ref_from_expr,
)
from app.services.formula_management.reporting_instructions import (
    DOC_TITLE,
    instructions_as_rows,
)

logger = logging.getLogger(__name__)

# 「编报说明」sheet 名（与前端说明弹窗同源标题呼应）。
INSTRUCTIONS_SHEET_NAME = "编报说明"
FORMULA_SHEET_NAME = "公式"

# 导出/导入列定义：(列头, 字段键)。列头用于 xlsx 头行匹配，字段键用于程序读写。
_COLUMNS: list[tuple[str, str]] = [
    ("页面键(page_key)", "page_key"),
    ("目标单元(target_cell)", "target_cell"),
    ("公式表达式(expression)", "expression"),
    ("公式类型(formula_type)", "formula_type"),
    ("引用(refs, JSON)", "refs"),
    ("最近计算时间(last_computed_at)", "last_computed_at"),
    ("说明(description)", "description"),
]
_HEADERS = [c[0] for c in _COLUMNS]
_FIELD_KEYS = [c[1] for c in _COLUMNS]

# 模板示例公式行（三类型各一，演示填法；与编报说明口径一致）。
_EXAMPLE_ROWS: list[list[str]] = [
    [
        "workpaper:D2",
        "C15",
        "TB('1122') + TB('1123')",
        "auto_calc",
        json.dumps([{"formula_ref": "TB('1122')"}, {"formula_ref": "TB('1123')"}], ensure_ascii=False),
        "",
        "应收账款 = 明细之和（计算回填示例）",
    ],
    [
        "report:cross_check",
        "report-cross-check-1",
        "ABS(ROW('assets_total') - (ROW('liabilities_total') + ROW('equity_total'))) <= 1",
        "logic_check",
        json.dumps(
            [
                {"formula_ref": "ROW('assets_total')"},
                {"formula_ref": "ROW('liabilities_total')"},
                {"formula_ref": "ROW('equity_total')"},
            ],
            ensure_ascii=False,
        ),
        "",
        "资产合计 = 负债合计 + 所有者权益合计（逻辑校验示例）",
    ],
    [
        "report:cross_check",
        "report-cross-check-6",
        "ABS(ROW('IS-018') - ROW('IS-017') * 0.25) > ROW('IS-017') * 0.05",
        "reasonability",
        json.dumps([{"formula_ref": "ROW('IS-017')"}, {"formula_ref": "ROW('IS-018')"}], ensure_ascii=False),
        "",
        "有效税率≈25%，偏离过大时提示核查（合理性提示示例）",
    ],
]


# ── 导入解析结果模型 ──────────────────────────────────────────────────────────
@dataclass
class ImportSkip:
    """一条被跳过的公式（悬空引用 / 无效）。"""

    row_number: int
    page_key: str
    target_cell: str
    reason: str
    dangling_refs: list[str] = field(default_factory=list)

    def to_dict(self) -> dict[str, Any]:
        return {
            "row_number": self.row_number,
            "page_key": self.page_key,
            "target_cell": self.target_cell,
            "reason": self.reason,
            "dangling_refs": list(self.dangling_refs),
        }


@dataclass
class ImportResult:
    """导入解析汇总：有效条目 + 跳过报告。"""

    imported: list[PresetEntry] = field(default_factory=list)
    skipped: list[ImportSkip] = field(default_factory=list)

    @property
    def imported_count(self) -> int:
        return len(self.imported)

    @property
    def skipped_count(self) -> int:
        return len(self.skipped)

    def to_dict(self) -> dict[str, Any]:
        return {
            "ok": True,
            "imported_count": self.imported_count,
            "skipped_count": self.skipped_count,
            "imported": [e.to_dict() for e in self.imported],
            "skipped": [s.to_dict() for s in self.skipped],
        }


# ── 编报说明 sheet 写入（Req 23.2，单一源） ───────────────────────────────────
def _write_instructions_sheet(ws) -> None:
    """把单一源编报说明写入工作簿首个 sheet（Req 23.2/23.6）。"""
    ws.column_dimensions["A"].width = 96
    first = True
    for row in instructions_as_rows():
        ws.append(row)
        cell = ws.cell(row=ws.max_row, column=1)
        cell.alignment = Alignment(wrap_text=True, vertical="top")
        text = row[0] if row else ""
        if first:
            cell.font = Font(bold=True, size=13)
            first = False
        elif text == DOC_TITLE or (text and text.startswith(("一、", "二、", "三、", "四、", "五、", "六、", "七、"))):
            cell.font = Font(bold=True, size=11)


def _write_formula_header(ws) -> None:
    ws.append(_HEADERS)
    for col_idx in range(1, len(_HEADERS) + 1):
        ws.cell(row=ws.max_row, column=col_idx).font = Font(bold=True)
        ws.column_dimensions[get_column_letter(col_idx)].width = 26
    ws.freeze_panes = "A2"


def _entry_to_row(entry: PresetEntry, *, last_computed_at: str | None = None) -> list[Any]:
    return [
        entry.page_key,
        entry.target_cell,
        entry.expression,
        entry.formula_type,
        json.dumps(entry.refs or [], ensure_ascii=False),
        last_computed_at or "",
        entry.description or "",
    ]


# ── 导出模板（Req 23.2） ──────────────────────────────────────────────────────
def build_template_workbook() -> Workbook:
    """构建导出模板工作簿：首区块=编报说明 + 公式 sheet（空白/示例行）。"""
    wb = Workbook()
    ws_doc = wb.active
    ws_doc.title = INSTRUCTIONS_SHEET_NAME
    _write_instructions_sheet(ws_doc)

    ws = wb.create_sheet(FORMULA_SHEET_NAME)
    _write_formula_header(ws)
    for example in _EXAMPLE_ROWS:
        ws.append(example)
    return wb


# ── 导出数据（Req 23.3） ──────────────────────────────────────────────────────
def build_data_workbook(
    entries: list[PresetEntry],
    *,
    last_computed_map: dict[tuple[str, str], str] | None = None,
) -> Workbook:
    """构建数据工作簿：首区块=编报说明 + 公式 sheet（已有公式）。

    Args:
        entries: 待导出的公式条目（预设库当前页面/模块条目）。
        last_computed_map: (page_key, target_cell) → 最近计算时间 ISO 串（可选）。
    """
    wb = Workbook()
    ws_doc = wb.active
    ws_doc.title = INSTRUCTIONS_SHEET_NAME
    _write_instructions_sheet(ws_doc)

    ws = wb.create_sheet(FORMULA_SHEET_NAME)
    _write_formula_header(ws)
    lcm = last_computed_map or {}
    for entry in entries:
        lc = lcm.get((entry.page_key, entry.target_cell))
        ws.append(_entry_to_row(entry, last_computed_at=lc))
    return wb


def workbook_to_bytes(wb: Workbook) -> bytes:
    buf = io.BytesIO()
    wb.save(buf)
    return buf.getvalue()


# ── 导入解析 + 逐条 full_resolve 校验（Req 23.4） ─────────────────────────────
def _parse_refs_cell(raw: Any) -> list[Any]:
    """解析 refs 列（JSON 数组）；非法/空返回空列表。"""
    if raw is None:
        return []
    if isinstance(raw, list):
        return raw
    text = str(raw).strip()
    if not text:
        return []
    try:
        parsed = json.loads(text)
        return parsed if isinstance(parsed, list) else []
    except (json.JSONDecodeError, TypeError):
        return []


def _normalize_ref_item(ref: Any) -> tuple[str | None, str | None]:
    """把一条 ref 归一化为 (formula_ref, addr_id)。"""
    if isinstance(ref, dict):
        return ref.get("formula_ref"), ref.get("addr_id")
    if isinstance(ref, str):
        return ref, None
    return None, None


def _read_header_row(ws) -> tuple[int, dict[str, int]] | None:
    """在公式 sheet 中定位表头行，返回 (header_row_index, {列头: 列下标})。"""
    for r_idx, row in enumerate(ws.iter_rows(min_row=1, max_row=5, values_only=True), start=1):
        cells = [("" if v is None else str(v).strip()) for v in row]
        # 命中：包含 page_key / target_cell / expression 三个关键列
        if _HEADERS[0] in cells and _HEADERS[1] in cells and _HEADERS[2] in cells:
            index = {h: i for i, h in enumerate(cells) if h}
            return r_idx, index
    return None


def _pick_formula_sheet(wb):
    """选择公式 sheet：优先名为「公式」，否则第一个含公式表头的 sheet。"""
    if FORMULA_SHEET_NAME in wb.sheetnames:
        ws = wb[FORMULA_SHEET_NAME]
        if _read_header_row(ws):
            return ws
    for ws in wb.worksheets:
        if _read_header_row(ws):
            return ws
    return None


async def parse_import_rows(
    content: bytes,
    *,
    db: AsyncSession | None = None,
    project_id: str | None = None,
) -> ImportResult:
    """解析上传 xlsx 并逐条经 ACNR full_resolve 校验引用（Req 23.4）。

    - 每条公式的引用逐个经 ``resolve_ref`` 解析：**任一引用悬空（found=False 且非
      fail-open）** → 整条公式报告并跳过（不入库）。
    - 基础设施故障（fail-open）不视为悬空，不跳过（Req 11.3）。
    - 无效公式类型 / 缺目标单元 → 跳过并报告。
    - refs 列留空时从表达式自动提取 formula_ref 再校验。

    Returns:
        ImportResult：有效条目 + 跳过报告。
    """
    result = ImportResult()
    try:
        wb = load_workbook(io.BytesIO(content), read_only=True, data_only=True)
    except Exception as exc:  # noqa: BLE001
        raise ValueError(f"无法解析 xlsx 文件: {exc}") from exc

    ws = _pick_formula_sheet(wb)
    if ws is None:
        wb.close()
        raise ValueError(
            f"未找到公式数据表（需包含列：{_HEADERS[0]} / {_HEADERS[1]} / {_HEADERS[2]}）"
        )

    header_info = _read_header_row(ws)
    assert header_info is not None
    header_row, col_index = header_info

    def _cell(row: tuple, header: str) -> Any:
        idx = col_index.get(header)
        if idx is None or idx >= len(row):
            return None
        return row[idx]

    row_number = header_row
    for row in ws.iter_rows(min_row=header_row + 1, values_only=True):
        row_number += 1
        if row is None or all(v is None for v in row):
            continue

        page_key = _safe_str(_cell(row, _HEADERS[0]))
        target_cell = _safe_str(_cell(row, _HEADERS[1]))
        expression = _safe_str(_cell(row, _HEADERS[2]))
        formula_type = _safe_str(_cell(row, _HEADERS[3]))
        refs_raw = _cell(row, _HEADERS[4])
        description = _safe_str(_cell(row, _HEADERS[6]))

        # 跳过完全空行 / 只填了说明的行
        if not page_key and not target_cell and not expression:
            continue

        # 基本字段校验
        if not page_key or not target_cell:
            result.skipped.append(
                ImportSkip(row_number, page_key, target_cell, "缺少 page_key 或 target_cell")
            )
            continue
        if formula_type not in VALID_FORMULA_TYPES:
            result.skipped.append(
                ImportSkip(
                    row_number,
                    page_key,
                    target_cell,
                    f"无效公式类型: {formula_type!r}（须为 {sorted(VALID_FORMULA_TYPES)} 之一）",
                )
            )
            continue

        refs = _parse_refs_cell(refs_raw)
        if not refs:
            # refs 留空时从表达式自动提取
            refs = _formula_ref_from_expr(expression)

        # 逐条引用经 ACNR full_resolve 校验
        dangling: list[str] = []
        for ref in refs:
            formula_ref, addr_id = _normalize_ref_item(ref)
            if not formula_ref and not addr_id:
                continue
            rr = await resolve_ref(
                formula_ref=formula_ref,
                addr_id=addr_id,
                project_id=project_id,
                db=db,
            )
            # fail-open（基础设施故障）不视为悬空，不跳过（Req 11.3）
            if not rr.found and not rr.fail_open:
                dangling.append(formula_ref or addr_id or "")

        if dangling:
            result.skipped.append(
                ImportSkip(
                    row_number,
                    page_key,
                    target_cell,
                    "引用悬空（无法经 ACNR 解析）",
                    dangling_refs=dangling,
                )
            )
            continue

        result.imported.append(
            PresetEntry(
                page_key=page_key,
                target_cell=target_cell,
                expression=expression,
                formula_type=formula_type,
                refs=refs,
                source="import",
                description=description,
            )
        )

    wb.close()
    return result


def _safe_str(val: Any) -> str:
    return "" if val is None else str(val).strip()

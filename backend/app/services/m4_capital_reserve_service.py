"""M4 资本公积 — 业务逻辑服务

提供功能：
- 导出模板 / 导出数据 / 导入数据（xlsx三级）
- 资本公积变动明细汇总
- J3股份支付权益结算联动核对
- M2外币折算差异联动

科目4002资本公积（**贷方/权益类！**）：期末=期初+贷方-借方
资本溢价+其他资本公积在贷方增加，转出（转增资本等）在借方减少。

Requirements: 6.1-6.2, 6.5
"""

from __future__ import annotations

import io
import json
import logging
from typing import Any
from uuid import uuid4

import sqlalchemy as sa
from openpyxl import Workbook, load_workbook
from openpyxl.styles import Alignment, Font, PatternFill
from sqlalchemy.ext.asyncio import AsyncSession

logger = logging.getLogger(__name__)

_M4_ACCOUNT_CODE = "4002"
_ROW_LIMIT = 500
_TOLERANCE = 0.01


# ═══════════════════════════════════════════════════════════════════════════════
# 导入导出配置
# ═══════════════════════════════════════════════════════════════════════════════

_SHEET_CONFIGS: dict[str, dict[str, Any]] = {
    "M4-2": {
        "title": "M4-2 明细表（资本溢价+其他资本公积）",
        "headers": [
            "类别", "来源项目", "期初余额", "本期增加(贷方)",
            "本期减少(借方)", "期末余额", "变动原因", "关联底稿",
            "核对状态", "备注",
        ],
        "fields": [
            "category", "source", "begin", "increase",
            "decrease", "end_balance", "reason", "linked_wp",
            "check_status", "remark",
        ],
    },
    "M4-4": {
        "title": "M4-4 资本公积检查表",
        "headers": [
            "检查项目", "检查内容", "检查结果", "差异金额",
            "说明", "结论", "备注",
        ],
        "fields": [
            "check_item", "check_content", "check_result", "diff_amount",
            "explanation", "conclusion", "remark",
        ],
    },
}

_HEADER_FONT = Font(bold=True, size=11)
_HEADER_FILL = PatternFill(start_color="4472C4", end_color="4472C4", fill_type="solid")
_HEADER_FONT_WHITE = Font(bold=True, size=11, color="FFFFFF")
_HEADER_ALIGN = Alignment(horizontal="center", vertical="center", wrap_text=True)


# ═══════════════════════════════════════════════════════════════════════════════
# 辅助函数
# ═══════════════════════════════════════════════════════════════════════════════


def _parse_num(v: Any) -> float:
    """安全解析数值，None/空/NaN→0.0."""
    if v is None or v == "":
        return 0.0
    try:
        f = float(v)
        return 0.0 if f != f else f
    except (ValueError, TypeError):
        return 0.0


async def _get_wp_context(wp_id: str, db: AsyncSession) -> dict[str, Any]:
    """从 wp_id 取关联的 project_id."""
    result = await db.execute(
        sa.text("""
            SELECT wp.project_id, p.audit_year
            FROM working_papers wp
            JOIN projects p ON p.id = wp.project_id
            WHERE wp.id = :wp_id
        """),
        {"wp_id": wp_id},
    )
    row = result.fetchone()
    if not row:
        raise ValueError(f"底稿不存在: {wp_id}")
    return {"project_id": row.project_id, "year": int(row.audit_year or 2025)}


def _create_sheet_with_headers(wb: Workbook, config: dict[str, Any]) -> None:
    """在工作簿中创建带格式列头的sheet."""
    ws = wb.create_sheet(title=config["title"])
    for col_idx, header in enumerate(config["headers"], 1):
        cell = ws.cell(row=1, column=col_idx, value=header)
        cell.font = _HEADER_FONT_WHITE
        cell.fill = _HEADER_FILL
        cell.alignment = _HEADER_ALIGN
    # 设置列宽
    for col_idx in range(1, len(config["headers"]) + 1):
        ws.column_dimensions[ws.cell(row=1, column=col_idx).column_letter].width = 16


# ═══════════════════════════════════════════════════════════════════════════════
# 导出模板
# ═══════════════════════════════════════════════════════════════════════════════


async def export_template(
    wp_id: str,
    db: AsyncSession,
    *,
    sheet: str | None = None,
) -> io.BytesIO:
    """导出空白xlsx模板（含列头格式）。

    Args:
        wp_id: 底稿ID
        db: 数据库会话
        sheet: 可选，指定导出单个sheet（M4-2/M4-4）
    """
    wb = Workbook()
    wb.remove(wb.active)  # type: ignore[arg-type]

    if sheet and sheet in _SHEET_CONFIGS:
        _create_sheet_with_headers(wb, _SHEET_CONFIGS[sheet])
    else:
        for cfg in _SHEET_CONFIGS.values():
            _create_sheet_with_headers(wb, cfg)

    buffer = io.BytesIO()
    wb.save(buffer)
    buffer.seek(0)
    return buffer


# ═══════════════════════════════════════════════════════════════════════════════
# 导出数据
# ═══════════════════════════════════════════════════════════════════════════════


async def export_data(
    wp_id: str,
    db: AsyncSession,
    *,
    sheet: str | None = None,
) -> io.BytesIO:
    """导出含当前数据的xlsx。

    从 checklist_responses 读取 M4 相关数据填入对应sheet。
    """
    wb = Workbook()
    wb.remove(wb.active)  # type: ignore[arg-type]

    # 读取 checklist_responses 数据
    responses = await _load_responses(wp_id, db)

    configs_to_export = (
        {sheet: _SHEET_CONFIGS[sheet]} if sheet and sheet in _SHEET_CONFIGS else _SHEET_CONFIGS
    )

    for sheet_key, cfg in configs_to_export.items():
        ws = wb.create_sheet(title=cfg["title"])
        # 写入列头
        for col_idx, header in enumerate(cfg["headers"], 1):
            cell = ws.cell(row=1, column=col_idx, value=header)
            cell.font = _HEADER_FONT_WHITE
            cell.fill = _HEADER_FILL
            cell.alignment = _HEADER_ALIGN

        # 写入数据行
        row_idx = 2
        prefix = f"M4-{sheet_key.split('-')[1]}-row-" if "-" in sheet_key else f"{sheet_key}-row-"
        for item_id, data in sorted(responses.items()):
            if not item_id.startswith(prefix):
                continue
            raw = data.get("conclusion", "")
            if not raw:
                continue
            try:
                row_data = json.loads(raw) if isinstance(raw, str) else raw
            except (json.JSONDecodeError, TypeError):
                continue
            if not isinstance(row_data, dict):
                continue

            for col_idx, field in enumerate(cfg["fields"], 1):
                val = row_data.get(field, "")
                ws.cell(row=row_idx, column=col_idx, value=val)
            row_idx += 1

    buffer = io.BytesIO()
    wb.save(buffer)
    buffer.seek(0)
    return buffer


# ═══════════════════════════════════════════════════════════════════════════════
# 导入数据
# ═══════════════════════════════════════════════════════════════════════════════


async def import_data(
    wp_id: str,
    content: bytes,
    db: AsyncSession,
    *,
    sheet: str | None = None,
) -> dict[str, Any]:
    """解析xlsx并写入checklist_responses。

    返回:
        { imported_count: int, warnings: list[str] }
    """
    warnings: list[str] = []
    imported_count = 0

    try:
        wb = load_workbook(io.BytesIO(content), read_only=True, data_only=True)
    except Exception as e:
        raise ValueError(f"无法解析xlsx文件: {e}")

    # 确定要导入的sheet配置
    configs_to_import: dict[str, dict[str, Any]] = {}
    if sheet and sheet in _SHEET_CONFIGS:
        configs_to_import = {sheet: _SHEET_CONFIGS[sheet]}
    else:
        configs_to_import = _SHEET_CONFIGS.copy()

    for sheet_key, cfg in configs_to_import.items():
        # 在工作簿中查找匹配的sheet
        ws = None
        for ws_name in wb.sheetnames:
            if cfg["title"] in ws_name or sheet_key in ws_name:
                ws = wb[ws_name]
                break

        if ws is None:
            warnings.append(f"未找到sheet: {cfg['title']}")
            continue

        # 验证列头
        actual_headers = [str(ws.cell(row=1, column=c).value or "").strip() for c in range(1, len(cfg["headers"]) + 1)]
        expected_headers = cfg["headers"]
        if actual_headers != expected_headers:
            # 允许部分匹配
            match_count = sum(1 for a, e in zip(actual_headers, expected_headers) if a == e)
            if match_count < len(expected_headers) * 0.6:
                warnings.append(f"Sheet {sheet_key} 列头不匹配: 期望 {expected_headers[:3]}..., 实际 {actual_headers[:3]}...")
                continue

        # 解析数据行
        sheet_suffix = sheet_key.split("-")[1] if "-" in sheet_key else sheet_key
        for row_idx in range(2, min(ws.max_row or 2, _ROW_LIMIT) + 1):
            row_data: dict[str, Any] = {}
            has_data = False
            for col_idx, field in enumerate(cfg["fields"], 1):
                val = ws.cell(row=row_idx, column=col_idx).value
                if val is not None and str(val).strip():
                    has_data = True
                row_data[field] = val

            if not has_data:
                continue

            item_id = f"M4-{sheet_suffix}-row-{row_idx - 1}"
            await db.execute(
                sa.text("""
                    INSERT INTO checklist_responses (id, workpaper_id, item_id, conclusion, status)
                    VALUES (:id, :wid, :iid, :conclusion, 'imported')
                    ON CONFLICT (workpaper_id, item_id)
                    DO UPDATE SET conclusion = :conclusion, status = 'imported'
                """),
                {
                    "id": str(uuid4()),
                    "wid": wp_id,
                    "iid": item_id,
                    "conclusion": json.dumps(row_data, ensure_ascii=False),
                },
            )
            imported_count += 1

    wb.close()
    return {"imported_count": imported_count, "warnings": warnings}


# ═══════════════════════════════════════════════════════════════════════════════
# 资本公积变动明细
# ═══════════════════════════════════════════════════════════════════════════════


async def get_reserve_changes(wp_id: str, db: AsyncSession) -> dict[str, Any]:
    """获取资本公积变动明细（从M4-2数据汇总）。

    分资本溢价/其他资本公积两类汇总。
    """
    responses = await _load_responses(wp_id, db)

    premium_changes: list[dict[str, Any]] = []
    other_changes: list[dict[str, Any]] = []

    for item_id, data in sorted(responses.items()):
        if not item_id.startswith("M4-2-row-"):
            continue
        raw = data.get("conclusion", "")
        if not raw:
            continue
        try:
            row_data = json.loads(raw) if isinstance(raw, str) else raw
        except (json.JSONDecodeError, TypeError):
            continue
        if not isinstance(row_data, dict):
            continue

        entry = {
            "source": row_data.get("source", ""),
            "begin": _parse_num(row_data.get("begin")),
            "increase": _parse_num(row_data.get("increase")),
            "decrease": _parse_num(row_data.get("decrease")),
            "end": _parse_num(row_data.get("end_balance")),
            "reason": row_data.get("reason", ""),
        }

        category = str(row_data.get("category", "")).strip()
        if "其他" in category:
            other_changes.append(entry)
        else:
            premium_changes.append(entry)

    # 汇总
    premium_total = sum(e["end"] for e in premium_changes)
    other_total = sum(e["end"] for e in other_changes)

    return {
        "premium_changes": premium_changes,
        "other_changes": other_changes,
        "summary": {
            "premium_total": premium_total,
            "other_total": other_total,
            "grand_total": premium_total + other_total,
        },
        "account_code": _M4_ACCOUNT_CODE,
    }


# ═══════════════════════════════════════════════════════════════════════════════
# J3 联动
# ═══════════════════════════════════════════════════════════════════════════════


async def get_j3_linkage(wp_id: str, db: AsyncSession) -> dict[str, Any]:
    """获取J3股份支付权益结算联动数据。

    从J3底稿查找等待期确认金额，与M4其他资本公积增加额核对。
    """
    ctx = await _get_wp_context(wp_id, db)
    project_id = ctx["project_id"]

    # ─── 从J3底稿取股份支付权益结算金额 ──────────────────────────────────
    j3_equity_settled: float = 0.0
    try:
        j3_result = await db.execute(
            sa.text("""
                SELECT cr.conclusion
                FROM checklist_responses cr
                JOIN working_papers wp ON wp.id = cr.workpaper_id
                JOIN wp_index wi ON wi.project_id = wp.project_id
                    AND wi.wp_code = 'J3'
                WHERE wp.project_id = :pid
                    AND cr.item_id LIKE 'J3-equity-settled%'
                LIMIT 1
            """),
            {"pid": str(project_id)},
        )
        j3_row = j3_result.fetchone()
        if j3_row and j3_row.conclusion:
            j3_equity_settled = _parse_num(j3_row.conclusion)
    except Exception as e:  # noqa: BLE001
        logger.warning("M4 J3联动查询失败: %s", e)

    # ─── 从M2底稿取外币折算差异 ──────────────────────────────────────────
    m2_fx_diff: float = 0.0
    try:
        m2_result = await db.execute(
            sa.text("""
                SELECT cr.conclusion
                FROM checklist_responses cr
                JOIN working_papers wp ON wp.id = cr.workpaper_id
                WHERE wp.project_id = :pid
                    AND cr.item_id LIKE 'M2-fx-diff%'
                LIMIT 1
            """),
            {"pid": str(project_id)},
        )
        m2_row = m2_result.fetchone()
        if m2_row and m2_row.conclusion:
            m2_fx_diff = _parse_num(m2_row.conclusion)
    except Exception as e:  # noqa: BLE001
        logger.warning("M4 M2联动查询失败: %s", e)

    # ─── 从M4自身取其他资本公积增加额 ────────────────────────────────────
    responses = await _load_responses(wp_id, db)
    m4_other_increase: float = 0.0
    for item_id, data in responses.items():
        if not item_id.startswith("M4-2-row-"):
            continue
        raw = data.get("conclusion", "")
        if not raw:
            continue
        try:
            row_data = json.loads(raw) if isinstance(raw, str) else raw
        except (json.JSONDecodeError, TypeError):
            continue
        if not isinstance(row_data, dict):
            continue
        category = str(row_data.get("category", "")).strip()
        if "其他" in category:
            m4_other_increase += _parse_num(row_data.get("increase"))

    diff = j3_equity_settled - m4_other_increase
    is_consistent = abs(diff) <= _TOLERANCE

    # 关联底稿引用
    cross_refs: list[str] = []
    if j3_equity_settled != 0:
        cross_refs.append("J3")
    if m2_fx_diff != 0:
        cross_refs.append("M2")

    return {
        "j3_equity_settled": j3_equity_settled,
        "m4_other_increase": m4_other_increase,
        "diff": round(diff, 2),
        "is_consistent": is_consistent,
        "m2_fx_diff": m2_fx_diff,
        "cross_refs": cross_refs,
    }


# ═══════════════════════════════════════════════════════════════════════════════
# 内部辅助
# ═══════════════════════════════════════════════════════════════════════════════


async def _load_responses(wp_id: str, db: AsyncSession) -> dict[str, dict[str, Any]]:
    """加载底稿的 checklist_responses."""
    result = await db.execute(
        sa.text(
            "SELECT item_id, conclusion, evidence, status "
            "FROM checklist_responses "
            "WHERE workpaper_id = :wid"
        ),
        {"wid": wp_id},
    )
    responses: dict[str, dict[str, Any]] = {}
    for r in result.fetchall():
        responses[r.item_id] = {
            "conclusion": r.conclusion,
            "evidence": r.evidence,
            "status": r.status,
        }
    return responses

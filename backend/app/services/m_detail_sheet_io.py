"""M 循环权益类明细表（M7-2 / M9-2）导入导出共用内核。

两张明细表在界面上的存储形态相同 ——「整表一个 JSON 键 + 逐行派生键 + 合计键」：

* 整表 ``{full_item_id}``：行数组，存在界面读取的那一列（M7 是 remark，M9 是 conclusion）；
* 逐行 ``{row_item_prefix}{i}-data`` / ``{row_item_prefix}{i}-auditedEnd``（i 从 0 起，remark）；
* 合计 ``{total_item_id}``（remark）。

本模块只放机制：模板 / 数据工作簿、读回、解析校验、整表替换写库。列定义、存储键与计算公式
由各科目 service 的 :class:`DetailSheetSpec` 给出；不认识的 sheet 由各 service 自己拒绝
（白名单归属科目，X-3 Deviation_Registry G6 也按 service 源码判定）。

SQL 只用 PostgreSQL 与 SQLite 共有的语法；本模块只执行语句、不提交，由路由统一 commit。
"""

from __future__ import annotations

import io
import json
import logging
import math
import re
import uuid
from dataclasses import dataclass
from datetime import date, datetime, time
from decimal import Decimal
from typing import Any, Callable, Literal, Mapping, Sequence

import sqlalchemy as sa
from fastapi import HTTPException
from openpyxl import Workbook, load_workbook
from openpyxl.styles import Alignment, Font, PatternFill
from openpyxl.utils import get_column_letter
from sqlalchemy.ext.asyncio import AsyncSession

logger = logging.getLogger(__name__)

#: 单次导入的数据行上限（与平台其余导入一致）
ROW_LIMIT = 500
#: 上传文件大小上限
MAX_FILE_BYTES = 10 * 1024 * 1024
#: 表头只在前若干行里找（允许文件顶部有标题行）
HEADER_SCAN_ROWS = 10
#: 报错时最多列出的问题条数
MAX_REPORTED_ERRORS = 10
#: 这些名称的行视为汇总行：导出时由系统生成，导入时跳过
TOTAL_ROW_NAMES: frozenset[str] = frozenset({"合计", "小计", "总计"})
GUIDE_SHEET_TITLE = "编制说明"
#: 与 Wp_Bound_Gate 拒绝时同一文案（不区分「不存在」与「无权访问」）
NOT_FOUND_DETAIL = "资源不存在或不可访问"

StorageColumn = Literal["conclusion", "remark"]
_STORAGE_COLUMNS: frozenset[str] = frozenset({"conclusion", "remark"})
#: 逐行派生键的后缀（与前端 `_triggerSave` 写的两种键一致）
_ROW_SUFFIXES: tuple[str, ...] = ("data", "auditedEnd")

_AMOUNT_FORMAT = "#,##0.00"
_RATIO_FORMAT = "0.00%"
_INPUT_HEADER_FILL = PatternFill(start_color="4472C4", end_color="4472C4", fill_type="solid")
_INPUT_HEADER_FONT = Font(bold=True, color="FFFFFF")
_COMPUTED_HEADER_FILL = PatternFill(start_color="D9E1F2", end_color="D9E1F2", fill_type="solid")
_COMPUTED_HEADER_FONT = Font(bold=True, color="1F3864")
_HEADER_ALIGN = Alignment(horizontal="center", vertical="center", wrap_text=True)
_BOLD = Font(bold=True)

_NUMBER_RE = re.compile(r"[+-]?(?:\d+(?:\.\d*)?|\.\d+)(?:[eE][+-]?\d+)?")
#: 会计表里单独一个短横线表示 0
_ZERO_MARKS: frozenset[str] = frozenset({"-", "--", "—", "－", "–"})


def norm_header(value: Any) -> str:
    """表头归一：去掉全部空白、全角括号转半角 ⇒「贷方（增加）」与「贷方(增加)」是同一列。"""
    if value is None:
        return ""
    return re.sub(r"\s+", "", str(value).replace("（", "(").replace("）", ")"))


def unknown_sheet_message(sheet: str, supported: Sequence[str]) -> str:
    """不认识的 sheet 的报错文案（与 X-3 共享实现同一格式）。"""
    return f"不支持的sheet: {sheet}。支持: {sorted(supported)}"


@dataclass(frozen=True)
class InputColumn:
    """可编辑列：导入读它、导出写它。``field`` 逐字照抄前端行对象的字段名。"""

    header: str
    field: str
    kind: Literal["text", "amount", "choice"] = "amount"
    #: choice 列：存储值 → 显示文字（导出写显示文字，导入两者都认）
    choices: Mapping[str, str] | None = None
    #: choice 列额外接受的写法 → 存储值（如审定表里的长名称）
    choice_aliases: Mapping[str, str] | None = None
    #: choice 列库里为空时按此值读（与界面 ``|| 默认值`` 一致）
    choice_default: str = ""


@dataclass(frozen=True)
class ComputedColumn:
    """公式列：只出现在「导出数据」里；导入时忽略，界面按同一公式重算。"""

    header: str
    field: str
    kind: Literal["amount", "ratio"] = "amount"


@dataclass(frozen=True)
class DetailSheetSpec:
    """一张明细表的列定义、存储键与计算公式。

    ``input_columns`` 的顺序就是前端写行 JSON 时的字段顺序（模板列序也取它），
    这样后端写出的 JSON 与界面自己保存的逐字节一致。
    """

    sheet_code: str
    sheet_title: str
    full_item_id: str
    full_storage: StorageColumn
    #: 整表键读不到时依次尝试的旧键（与界面恢复顺序一致）
    fallback_sources: tuple[tuple[str, StorageColumn], ...]
    row_item_prefix: str
    total_item_id: str
    #: 导入生成的行 key 前缀（界面恢复时原样保留 key）
    row_key_prefix: str
    input_columns: tuple[InputColumn, ...]
    computed_columns: tuple[ComputedColumn, ...]
    #: 单行 → 公式列（必须含 ``auditedEnd``），逐式照抄前端（含运算顺序）
    compute: Callable[[Mapping[str, Any]], dict[str, float]]
    #: 全部行的公式列 → 审定期末合计（照抄前端 totals 的算法）
    audited_end_total: Callable[[Sequence[Mapping[str, float]]], float]
    guidance: tuple[str, ...] = ()

    def __post_init__(self) -> None:
        problems: list[str] = []
        if not self.input_columns or self.input_columns[0].kind != "text":
            problems.append("第一列必须是每行必填的名称文本列")
        storages = {self.full_storage, *(column for _item, column in self.fallback_sources)}
        if not storages <= _STORAGE_COLUMNS:
            problems.append(f"存储列只能是 {sorted(_STORAGE_COLUMNS)}")
        columns = (*self.input_columns, *self.computed_columns)
        headers = [norm_header(col.header) for col in columns]
        if len(set(headers)) != len(headers):
            problems.append("列名归一后有重复")
        fields = [col.field for col in columns]
        if len(set(fields)) != len(fields) or "key" in fields:
            problems.append("字段名重复或占用了 key")
        problems.extend(
            f"选择列「{col.header}」缺少 choices"
            for col in self.input_columns
            if col.kind == "choice" and not col.choices
        )
        if problems:
            raise ValueError(f"{self.sheet_code} 明细表规格无效：" + "；".join(problems))

    @property
    def key_column(self) -> InputColumn:
        """每行必填的名称列（项目名称 / OCI项目）。"""
        return self.input_columns[0]

    def blank_row(self) -> dict[str, Any]:
        """全部输入字段取初值的行（金额 0、文本空、选择列取默认值）。"""
        row: dict[str, Any] = {}
        for col in self.input_columns:
            if col.kind == "amount":
                row[col.field] = 0.0
            elif col.kind == "choice":
                row[col.field] = col.choice_default
            else:
                row[col.field] = ""
        return row


# ── 数字与文本 ──────────────────────────────────────────────────────────────


def js_number(value: float) -> str:
    """与前端 ``String(number)`` 同形：整数不带小数点，``1e-6 ≤ |x| < 1e21`` 用定点写法。"""
    num = float(value)
    if not math.isfinite(num):
        raise ValueError(f"非有限数字不能写入: {value!r}")
    if num == 0:
        return "0"
    text = repr(num)
    if "e" in text:
        if 1e-6 <= abs(num) < 1e21:
            text = format(Decimal(text), "f")
        else:
            mantissa, _, exponent = text.partition("e")
            sign = "-" if exponent.startswith("-") else "+"
            text = f"{mantissa}e{sign}{exponent.lstrip('+-').lstrip('0') or '0'}"
    if text.endswith(".0"):
        text = text[:-2]
    return text


def _json_value(value: Any) -> Any:
    """``JSON.stringify`` 写整数不带 ``.0``；非有限数字前端会写成 null，这里直接拒绝。"""
    if isinstance(value, float):
        if not math.isfinite(value):
            raise ValueError(f"非有限数字不能写入: {value!r}")
        if value == 0:
            return 0
        if value.is_integer() and abs(value) < 1e21:
            return int(value)
    return value


def _dumps(payload: Mapping[str, Any] | Sequence[Mapping[str, Any]]) -> str:
    """与前端 ``JSON.stringify`` 同形（紧凑分隔、整数不带 ``.0``、中文不转义）。"""
    if isinstance(payload, Mapping):
        normalized: Any = {key: _json_value(value) for key, value in payload.items()}
    else:
        normalized = [{key: _json_value(value) for key, value in row.items()} for row in payload]
    return json.dumps(normalized, ensure_ascii=False, separators=(",", ":"))


def lenient_number(value: Any) -> float:
    """读库用的宽松数值，与前端 ``safe()`` 同口径：缺失 / 非数字 / 非有限 ⇒ 0。"""
    if isinstance(value, bool):
        return 1.0 if value else 0.0
    if isinstance(value, (int, float)):
        num = float(value)
    elif isinstance(value, str):
        text = value.strip()
        if not text:
            return 0.0
        try:
            num = float(text)
        except ValueError:
            return 0.0
    else:
        return 0.0
    return num if math.isfinite(num) else 0.0


def parse_amount(value: Any) -> float:
    """导入用的金额解析；认不出来抛 ``ValueError``（调用方补行列号）。

    空白 ⇒ 0（与界面新增行的初值一致）；接受千分位、货币符号、``(123)`` 表示负数、
    单独的短横线表示 0。布尔值、日期一律拒绝 —— 填在金额列里多半是填错了列。
    """
    if value is None:
        return 0.0
    if isinstance(value, (bool, datetime, date, time)):
        raise ValueError(str(value))
    if isinstance(value, (int, float)):
        num = float(value)
    else:
        text = str(value).strip()
        if not text or text in _ZERO_MARKS:
            return 0.0
        negative = len(text) > 1 and (text[0], text[-1]) in {("(", ")"), ("（", "）")}
        if negative:
            text = text[1:-1].strip()
        text = re.sub(r"[,，\s¥￥]", "", text)
        if not _NUMBER_RE.fullmatch(text):
            raise ValueError(str(value))
        num = -float(text) if negative else float(text)
    if not math.isfinite(num):
        raise ValueError(str(value))
    return num


def _cell_text(value: Any) -> str:
    if value is None:
        return ""
    if isinstance(value, float) and value.is_integer():
        return str(int(value))
    return str(value).strip()


def _is_blank(value: Any) -> bool:
    return value is None or (isinstance(value, str) and not value.strip())


def _preview(value: Any) -> str:
    text = _cell_text(value)
    return text[:20] if text else "（空）"


def _parse_choice(col: InputColumn, value: Any) -> str:
    """选择列：存储值、显示文字、别名都认（去空白、括号归一、英文不分大小写）。"""
    def _key(text: Any) -> str:
        return norm_header(text).lower()

    lookup: dict[str, str] = {}
    for stored, label in (col.choices or {}).items():
        lookup[_key(stored)] = stored
        lookup[_key(label)] = stored
    for alias, stored in (col.choice_aliases or {}).items():
        lookup[_key(alias)] = stored
    found = lookup.get(_key(_cell_text(value)))
    if found is None:
        raise ValueError(str(value))
    return found


def normalize_stored_row(spec: DetailSheetSpec, raw: Mapping[str, Any]) -> dict[str, Any]:
    """库里的一行 → 只含 key 与输入字段的行（与界面恢复时 ``_toDetailRow`` 同口径）。"""
    row: dict[str, Any] = {"key": _cell_text(raw.get("key"))}
    for col in spec.input_columns:
        value = raw.get(col.field)
        if col.kind == "amount":
            row[col.field] = lenient_number(value)
        elif col.kind == "choice":
            row[col.field] = _cell_text(value) or col.choice_default
        else:
            row[col.field] = _cell_text(value)
    return row


# ── 底稿与读库 ──────────────────────────────────────────────────────────────


def require_wp_uuid(wp_id: Any) -> uuid.UUID:
    """路径里的底稿 id 必须是 UUID —— 否则 gate 不做判定（no-op），查询会在 PG 上报类型错误。"""
    try:
        return uuid.UUID(str(wp_id))
    except (ValueError, TypeError, AttributeError) as exc:
        raise HTTPException(status_code=404, detail=NOT_FOUND_DETAIL) from exc


async def require_project_id(db: AsyncSession, wp_id: Any) -> str:
    """底稿所属项目（``checklist_responses.project_id`` 非空）；底稿不存在或已删除 ⇒ 404。"""
    wp_uuid = require_wp_uuid(wp_id)
    result = await db.execute(
        sa.text("SELECT project_id FROM working_paper WHERE id = :wp_id AND is_deleted = false"),
        {"wp_id": str(wp_uuid)},
    )
    project_id = result.scalar_one_or_none()
    if project_id is None:
        raise HTTPException(status_code=404, detail=NOT_FOUND_DETAIL)
    return str(project_id)


def like_prefix(prefix: str) -> str:
    """LIKE 前缀匹配模式：``%`` ``_`` ``\\`` 先转义，否则会多捞别的键。"""
    return re.sub(r"([\\%_])", r"\\\1", prefix) + "%"


async def _fetch_columns(
    db: AsyncSession, wp_uuid: uuid.UUID, item_ids: Sequence[str]
) -> dict[str, dict[str, Any]]:
    stmt = sa.text(
        "SELECT item_id, conclusion, remark FROM checklist_responses "
        "WHERE wp_id = :wp_id AND item_id IN :item_ids"
    ).bindparams(sa.bindparam("item_ids", expanding=True))
    result = await db.execute(stmt, {"wp_id": str(wp_uuid), "item_ids": list(item_ids)})
    return {row.item_id: {"conclusion": row.conclusion, "remark": row.remark} for row in result.fetchall()}


async def load_rows(
    db: AsyncSession, wp_id: Any, spec: DetailSheetSpec
) -> tuple[list[dict[str, Any]], list[str]]:
    """按界面恢复的顺序读回明细行：整表键 → 旧键；空数组 / 坏 JSON 落到下一个来源。

    返回 ``(rows, warnings)``：坏数据不静默吞掉，写进导出文件的「编制说明」。
    """
    wp_uuid = require_wp_uuid(wp_id)
    sources = ((spec.full_item_id, spec.full_storage), *spec.fallback_sources)
    stored = await _fetch_columns(db, wp_uuid, [item_id for item_id, _column in sources])
    warnings: list[str] = []
    for item_id, column in sources:
        raw = (stored.get(item_id) or {}).get(column)
        if raw is None or not str(raw).strip():
            continue
        try:
            parsed = json.loads(raw)
        except (json.JSONDecodeError, TypeError) as exc:
            logger.error("%s 明细读回：%s.%s 不是有效 JSON（wp=%s）：%s", spec.sheet_code, item_id, column, wp_uuid, exc)
            warnings.append(f"库中「{item_id}」的内容不是有效的 JSON，已跳过")
            continue
        if not isinstance(parsed, list):
            warnings.append(f"库中「{item_id}」不是行列表，已跳过")
            continue
        if not parsed:
            continue
        rows = [normalize_stored_row(spec, item) for item in parsed if isinstance(item, Mapping)]
        if len(rows) != len(parsed):
            warnings.append(f"库中「{item_id}」有 {len(parsed) - len(rows)} 个元素不是行对象，已跳过")
        return rows, warnings
    return [], warnings


# ── 工作簿 ──────────────────────────────────────────────────────────────────


def _new_workbook() -> Workbook:
    wb = Workbook()
    if wb.active is not None:
        wb.remove(wb.active)
    return wb


def _to_buffer(wb: Workbook) -> io.BytesIO:
    buffer = io.BytesIO()
    wb.save(buffer)
    buffer.seek(0)
    return buffer


def _write_values(
    ws: Any,
    row_no: int,
    columns: Sequence[InputColumn | ComputedColumn],
    values: Mapping[str, Any],
    *,
    bold: bool = False,
) -> None:
    for idx, col in enumerate(columns, start=1):
        value = values.get(col.field)
        if isinstance(col, InputColumn) and col.kind == "choice" and value:
            value = (col.choices or {}).get(str(value), value)
        elif isinstance(col, ComputedColumn) and isinstance(value, float):
            value = round(value, 6 if col.kind == "ratio" else 2)
        cell = ws.cell(row=row_no, column=idx, value=None if value == "" else value)
        if col.kind == "amount":
            cell.number_format = _AMOUNT_FORMAT
        elif col.kind == "ratio":
            cell.number_format = _RATIO_FORMAT
        if bold:
            cell.font = _BOLD


def _write_sheet(wb: Workbook, spec: DetailSheetSpec, rows: Sequence[Mapping[str, Any]] | None) -> None:
    """``rows is None`` ⇒ 模板（只有输入列表头）；否则写数据行、公式列与合计行。"""
    ws = wb.create_sheet(title=spec.sheet_title)
    columns: list[InputColumn | ComputedColumn] = list(spec.input_columns)
    if rows is not None:
        columns.extend(spec.computed_columns)
    for idx, col in enumerate(columns, start=1):
        cell = ws.cell(row=1, column=idx, value=col.header)
        computed = isinstance(col, ComputedColumn)
        cell.font = _COMPUTED_HEADER_FONT if computed else _INPUT_HEADER_FONT
        cell.fill = _COMPUTED_HEADER_FILL if computed else _INPUT_HEADER_FILL
        cell.alignment = _HEADER_ALIGN
        ws.column_dimensions[get_column_letter(idx)].width = 28 if idx == 1 else 15
    ws.freeze_panes = "B2"
    if not rows:
        return
    amount_fields = [col.field for col in spec.input_columns if col.kind == "amount"]
    totals = dict.fromkeys(amount_fields, 0.0)
    for offset, row in enumerate(rows):
        _write_values(ws, offset + 2, columns, {**row, **spec.compute(row)})
        for field in amount_fields:
            totals[field] += float(row[field])
    total_row = {**spec.blank_row(), **totals}
    total_values = {**total_row, **spec.compute(total_row), spec.key_column.field: "合计"}
    for col in spec.input_columns:
        if col.kind != "amount" and col is not spec.key_column:
            total_values[col.field] = ""
    _write_values(ws, len(rows) + 2, columns, total_values, bold=True)


def _write_guide(wb: Workbook, specs: Sequence[DetailSheetSpec], warnings: Sequence[str]) -> None:
    ws = wb.create_sheet(title=GUIDE_SHEET_TITLE)
    lines: list[str] = []
    for spec in specs:
        lines += [
            f"{spec.sheet_title} 导入导出说明",
            f"1. 第 1 行是表头，列名请勿修改；「{spec.key_column.header}」每行必填。",
            "2. 金额列填数字，可带千分位；空白按 0 处理，(123) 表示负数。",
            "3. 「合计」「小计」行由系统按明细计算，导入时自动跳过。",
            f"4. 导入为整表替换：文件中的行会替换该明细表现有的全部行，单次最多 {ROW_LIMIT} 行。",
            "5. 浅色表头的列是公式列（只在「导出数据」中出现），导入时忽略，界面按同一公式重算。",
        ]
        lines += [f"   · {rule}" for rule in spec.guidance]
        lines.append("")
    if warnings:
        lines.append("导出提示：")
        lines += [f"   · {warning}" for warning in warnings]
    for line in lines:
        ws.append([line])
    ws["A1"].font = _BOLD
    ws.column_dimensions["A"].width = 110


def export_template(specs: Sequence[DetailSheetSpec]) -> io.BytesIO:
    """空白模板：每张明细表一页（只有输入列表头）+「编制说明」页。"""
    wb = _new_workbook()
    for spec in specs:
        _write_sheet(wb, spec, None)
    _write_guide(wb, specs, ())
    return _to_buffer(wb)


async def export_data(db: AsyncSession, wp_id: Any, specs: Sequence[DetailSheetSpec]) -> io.BytesIO:
    """当前数据：输入列 + 公式列 + 合计行；读回问题写进「编制说明」页。"""
    wp_uuid = require_wp_uuid(wp_id)
    wb = _new_workbook()
    warnings: list[str] = []
    for spec in specs:
        rows, sheet_warnings = await load_rows(db, wp_uuid, spec)
        _write_sheet(wb, spec, rows)
        warnings += sheet_warnings
        if not rows:
            warnings.append(f"{spec.sheet_title} 暂无已保存的明细行")
    _write_guide(wb, specs, warnings)
    return _to_buffer(wb)


# ── 导入解析 ────────────────────────────────────────────────────────────────


@dataclass(frozen=True)
class ParsedSheet:
    rows: list[dict[str, Any]]
    warnings: list[str]


def parse_workbook(content: bytes, spec: DetailSheetSpec) -> ParsedSheet:
    """解析上传的工作簿；任何问题都整批拒绝（``ValueError``，路由转 400），不做部分导入。"""
    if not content:
        raise ValueError("上传的文件为空")
    if len(content) > MAX_FILE_BYTES:
        raise ValueError("文件大小不能超过 10MB")
    try:
        wb = load_workbook(io.BytesIO(content), read_only=True, data_only=True)
    except Exception as exc:  # noqa: BLE001 — openpyxl 对坏文件抛 BadZipFile / InvalidFileException / KeyError 等多种异常
        logger.warning("%s 导入：无法解析上传文件：%s", spec.sheet_code, exc)
        raise ValueError("无法解析文件，请上传 .xlsx 工作簿（可先「导出模板」再填写）") from exc
    try:
        ws, header_row, headers = _locate_table(wb, spec)
        positions, warnings = _map_headers(spec, headers)
        rows = _read_rows(spec, ws, header_row, positions, warnings)
    finally:
        wb.close()
    return ParsedSheet(rows, warnings)


def _locate_table(wb: Any, spec: DetailSheetSpec) -> tuple[Any, int, list[str]]:
    """标题含 sheet 码的工作表优先；前若干行内含名称列表头的那一行即表头。"""
    key = norm_header(spec.key_column.header)
    sheets = sorted(wb.worksheets, key=lambda ws: spec.sheet_code not in str(ws.title))
    for ws in sheets:
        if hasattr(ws, "reset_dimensions"):
            ws.reset_dimensions()  # 只读模式按文件里记录的范围读，记录不准时会漏行
        for row_no, values in enumerate(
            ws.iter_rows(min_row=1, max_row=HEADER_SCAN_ROWS, values_only=True), start=1
        ):
            headers = [norm_header(value) for value in values]
            if key in headers:
                return ws, row_no, headers
    raise ValueError(
        f"未找到「{spec.sheet_title}」的表头（应包含「{spec.key_column.header}」列），"
        "请使用「导出模板」生成的文件填写"
    )


def _map_headers(spec: DetailSheetSpec, headers: Sequence[str]) -> tuple[dict[str, int], list[str]]:
    inputs = {norm_header(col.header): col for col in spec.input_columns}
    computed = {norm_header(col.header) for col in spec.computed_columns}
    positions: dict[str, int] = {}
    duplicated: list[str] = []
    unknown: list[str] = []
    for idx, header in enumerate(headers):
        if not header or header in computed:
            continue
        col = inputs.get(header)
        if col is None:
            unknown.append(header)
        elif col.field in positions:
            duplicated.append(col.header)
        else:
            positions[col.field] = idx
    if duplicated:
        raise ValueError("表头重复：" + "、".join(dict.fromkeys(duplicated)))
    missing = [col.header for col in spec.input_columns if col.field not in positions]
    if missing:
        raise ValueError("缺少列：" + "、".join(missing) + "（请使用「导出模板」生成的文件）")
    warnings = ["已忽略无法识别的列：" + "、".join(dict.fromkeys(unknown))] if unknown else []
    return positions, warnings


def _read_rows(
    spec: DetailSheetSpec,
    ws: Any,
    header_row: int,
    positions: Mapping[str, int],
    warnings: list[str],
) -> list[dict[str, Any]]:
    key_field = spec.key_column.field
    rows: list[dict[str, Any]] = []
    errors: list[str] = []
    skipped_totals = 0
    for offset, values in enumerate(ws.iter_rows(min_row=header_row + 1, values_only=True)):
        excel_row = header_row + 1 + offset
        cells = {field: (values[idx] if idx < len(values) else None) for field, idx in positions.items()}
        if all(_is_blank(value) for value in cells.values()):
            continue
        if _cell_text(cells[key_field]) in TOTAL_ROW_NAMES:
            skipped_totals += 1
            continue
        if len(rows) >= ROW_LIMIT:
            raise ValueError(f"数据行超过 {ROW_LIMIT} 行上限，请拆分后分批导入")
        row, row_errors = _parse_row(spec, cells, excel_row)
        errors += row_errors
        rows.append(row)
    if errors:
        shown = "；".join(errors[:MAX_REPORTED_ERRORS])
        more = f"……（另有 {len(errors) - MAX_REPORTED_ERRORS} 处）" if len(errors) > MAX_REPORTED_ERRORS else ""
        raise ValueError(f"导入失败，共 {len(errors)} 处问题，未写入任何数据：{shown}{more}")
    if not rows:
        raise ValueError("文件中没有可导入的数据行（表头下方为空，或只有合计行）")
    if skipped_totals:
        warnings.append(f"已跳过合计 / 小计行 {skipped_totals} 行（合计由系统按明细重算）")
    return rows


def _parse_row(
    spec: DetailSheetSpec, cells: Mapping[str, Any], excel_row: int
) -> tuple[dict[str, Any], list[str]]:
    row: dict[str, Any] = {}
    errors: list[str] = []
    for col in spec.input_columns:
        raw = cells.get(col.field)
        if col.kind == "amount":
            try:
                row[col.field] = parse_amount(raw)
            except ValueError:
                errors.append(f"第{excel_row}行「{col.header}」不是有效数字：{_preview(raw)}")
                row[col.field] = 0.0
        elif col.kind == "choice":
            try:
                row[col.field] = _parse_choice(col, raw)
            except ValueError:
                options = "或".join(f"「{label}」" for label in (col.choices or {}).values())
                errors.append(f"第{excel_row}行「{col.header}」应填{options}，实际为：{_preview(raw)}")
                row[col.field] = col.choice_default
        else:
            row[col.field] = _cell_text(raw)
    if not row[spec.key_column.field]:
        errors.insert(0, f"第{excel_row}行「{spec.key_column.header}」为空")
    return row, errors


# ── 写库 ────────────────────────────────────────────────────────────────────


async def _stale_row_item_ids(
    db: AsyncSession, wp_uuid: uuid.UUID, spec: DetailSheetSpec, *, keep: int
) -> list[str]:
    """本次行数以外、仍有值的逐行派生键（行号 ≥ keep）。只认两种后缀，不碰同前缀的其他键。"""
    result = await db.execute(
        sa.text(
            "SELECT item_id, remark FROM checklist_responses "
            "WHERE wp_id = :wp_id AND item_id LIKE :pattern ESCAPE '\\'"
        ),
        {"wp_id": str(wp_uuid), "pattern": like_prefix(spec.row_item_prefix)},
    )
    shape = re.compile(rf"{re.escape(spec.row_item_prefix)}(\d+)-(?:{'|'.join(_ROW_SUFFIXES)})")
    stale: list[str] = []
    for item_id, remark in result.fetchall():
        match = shape.fullmatch(item_id)
        if match and int(match.group(1)) >= keep and remark is not None:
            stale.append(item_id)
    return sorted(stale)


async def _upsert(
    db: AsyncSession,
    wp_uuid: uuid.UUID,
    project_id: str,
    writes: Mapping[str, tuple[StorageColumn, str | None]],
    *,
    updated_by: Any,
) -> None:
    """按存储列分组批量 upsert；冲突时只改本键的存储列，另一列（界面可能存了别的）保持不动。"""
    actor = None if updated_by is None else str(updated_by)
    grouped: dict[str, list[dict[str, Any]]] = {}
    for item_id, (column, value) in writes.items():
        grouped.setdefault(column, []).append(
            {
                "id": str(uuid.uuid4()),
                "project_id": project_id,
                "wp_id": str(wp_uuid),
                "item_id": item_id,
                "value": value,
                "updated_by": actor,
            }
        )
    for column, params in grouped.items():
        if column not in _STORAGE_COLUMNS:
            raise ValueError(f"非法存储列: {column}")
        set_actor = ", updated_by = excluded.updated_by" if actor is not None else ""
        await db.execute(
            sa.text(
                f"INSERT INTO checklist_responses (id, project_id, wp_id, item_id, {column}, "  # noqa: S608 — 列名已白名单校验
                "updated_by, created_at, updated_at) "
                "VALUES (:id, :project_id, :wp_id, :item_id, :value, :updated_by, "
                "CURRENT_TIMESTAMP, CURRENT_TIMESTAMP) "
                f"ON CONFLICT (wp_id, item_id) DO UPDATE SET {column} = excluded.{column}"
                f"{set_actor}, updated_at = CURRENT_TIMESTAMP"
            ),
            params,
        )


async def replace_rows(
    db: AsyncSession,
    wp_id: Any,
    spec: DetailSheetSpec,
    rows: Sequence[Mapping[str, Any]],
    *,
    updated_by: Any = None,
) -> list[dict[str, Any]]:
    """整表替换：整表键 + 逐行派生键 + 合计键一次写入，超出行数的逐行键置空。

    返回实际落库的行（含生成的 key）。只执行语句不提交。
    """
    wp_uuid = require_wp_uuid(wp_id)
    project_id = await require_project_id(db, wp_uuid)
    batch = uuid.uuid4().hex[:8]
    stored = [
        {"key": f"{spec.row_key_prefix}-{i}-{batch}", **{col.field: row[col.field] for col in spec.input_columns}}
        for i, row in enumerate(rows)
    ]
    computed = [spec.compute(row) for row in stored]
    writes: dict[str, tuple[StorageColumn, str | None]] = {
        spec.full_item_id: (spec.full_storage, _dumps(stored)),
    }
    for i, (row, calc) in enumerate(zip(stored, computed)):
        writes[f"{spec.row_item_prefix}{i}-data"] = ("remark", _dumps(row))
        writes[f"{spec.row_item_prefix}{i}-auditedEnd"] = ("remark", js_number(calc["auditedEnd"]))
    writes[spec.total_item_id] = ("remark", js_number(spec.audited_end_total(computed)))
    for item_id in await _stale_row_item_ids(db, wp_uuid, spec, keep=len(stored)):
        writes[item_id] = ("remark", None)
    await _upsert(db, wp_uuid, project_id, writes, updated_by=updated_by)
    return stored


async def import_sheet(
    db: AsyncSession,
    wp_id: Any,
    content: bytes,
    spec: DetailSheetSpec,
    *,
    updated_by: Any = None,
) -> dict[str, Any]:
    """解析 + 整表替换。返回 ``imported_count`` / ``warnings``（另给 ``warning`` 拼接串供界面提示）。"""
    wp_uuid = require_wp_uuid(wp_id)
    parsed = parse_workbook(content, spec)
    stored = await replace_rows(db, wp_uuid, spec, parsed.rows, updated_by=updated_by)
    logger.info("%s 明细导入：wp=%s 行数=%d 提示=%d", spec.sheet_code, wp_uuid, len(stored), len(parsed.warnings))
    result: dict[str, Any] = {
        "ok": True,
        "sheet": spec.sheet_code,
        "imported_count": len(stored),
        "warnings": list(parsed.warnings),
    }
    if parsed.warnings:
        result["warning"] = "；".join(parsed.warnings)
    return result

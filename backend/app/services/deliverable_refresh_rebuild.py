"""交付件增量刷新的「按模板重建块 + 交付过滤」适配（spec ``deliverable-trio-generation-accuracy`` §16b）。

为什么需要（2026-10-17 落实此前延后的可选任务 16b）
--------------------------------------------------
``deliverable_refresh_service._replace_block_content`` 原来只把库内 ``text_content`` 按编辑器口径转 HTML
插回块内（纯换正文文字），**不重渲染库内表格、不接交付过滤**（design §2.9「未覆盖」如实记录）。而本 spec
的全量生成（``note_template_export_fill.fill_template_blocks`` → ``note_template_block_fill.rebuild_block``）
产出的交付件是：标题 + 库内有数表（含表题 / 子标题、克隆 / 回退渲染）+ **交付过滤后**的正文，指引 / 示例 /
占位 / 大模型草稿均被剔除。两条路径结构不一致：刷新一次就把过滤掉的大模型草稿、空样例表又带回来、库内表
变动也不反映。

本模块把增量刷新对齐到同一套重建逻辑：
- **可重建块**（块内有标题 / 已渲染表 / 正文，``rebuild_block`` 能解析出 layout）→ 调 ``rebuild_block``
  重渲染库内表 + 过滤正文；
- **纯文字块回退**（无可识别结构）→ 至少接入交付过滤（``make_block_filter``）再渲染正文。

与 ``rebuild_block`` 的 block 结构适配（关键）
--------------------------------------------
``rebuild_block`` 约定 ``block.elements[0]`` 是开标记、``[-1]`` 是闭标记（二者不删、作插入锚点），
内容是 ``elements[1:-1]``，删除走 ``open_el.getnext() … close_el``、插入走 ``open_el.addnext``
—— 要求开闭标记与内容**同父且相邻**。

- **marker 模式**（``##SECTION##`` 标记仍在，模板 / 中间态文档）：原块 ``.elements`` 恰好满足，直接传。
- **anchor 模式**（交付件：标记已清、锚点在，内容可能被包进 ``w:sdt`` 内容控件）：``AnchorBlock.elements``
  是**纯内容**（不含 bookmark），且 ``start_el`` 在 body 级、内容在 ``w:sdtContent`` 内 ⇒ 不能拿
  ``start_el`` 当兄弟锚点。改为用 ``anchor_block_content_host`` 取内容宿主（body 或 ``w:sdtContent``），
  在宿主内首尾插两个**哨兵空段**作 open / close，重建后删哨兵。哨兵与内容同父、相邻，满足约定。

lxml ``id()`` 复用坑：哨兵全程持有元素引用删除，不按裸 ``id()`` 判定（steering / 本 spec AC-39）。
"""

from __future__ import annotations

from dataclasses import dataclass
from typing import Any, Callable

from docx.oxml import OxmlElement
from docx.oxml.ns import qn

__all__ = ["RebuildOutcome", "rebuild_refresh_block"]

_W_P = qn("w:p")
_W_TBL = qn("w:tbl")


@dataclass
class _AdaptedBlock:
    """喂给 ``rebuild_block`` 的鸭子类型块：``elements[0]`` / ``[-1]`` 是哨兵或原标记。"""

    section_code: str
    elements: list[Any]
    start_el: Any
    open_el: Any
    close_el: Any


@dataclass
class RebuildOutcome:
    """重建结果（供调用方记录 / 断言）。"""

    rebuilt: bool  # True = 走了 rebuild_block（模板重建）；False = 纯文字回退过滤渲染
    tables_rendered: int = 0
    tables_cloned: int = 0
    dropped_lines: int = 0
    dropped_tables: int = 0
    placeholders_left: int = 0


def _content_host(block: Any, locate_mode: str) -> tuple[Any, list[Any]]:
    """内容宿主与宿主内当前内容元素（``w:p`` / ``w:tbl``）。

    anchor 模式下内容可能在 ``w:sdtContent`` 内；marker 模式下内容是块内标记之间的元素。
    """
    if locate_mode == "anchor":
        from app.services.section_anchor_utils import anchor_block_content_host

        return anchor_block_content_host(block)
    # marker：开闭标记之间的内容
    elements = list(getattr(block, "elements", None) or [])
    inner = elements[1:-1] if len(elements) >= 2 else []
    host = block.open_el.getparent() if elements else None
    return host, inner


def _adapt(block: Any, locate_mode: str) -> tuple[_AdaptedBlock, Callable[[], None]]:
    """把定位块转成 ``rebuild_block`` 可消费的块；返回 ``(块, 清理哨兵的函数)``。

    marker 模式直接复用原标记（``elements[0]`` / ``[-1]``），无哨兵；
    anchor 模式在内容宿主内首尾插哨兵空段，内容元素在二者之间。
    """
    if locate_mode == "marker":
        elements = list(block.elements)
        adapted = _AdaptedBlock(
            section_code=block.section_code,
            elements=elements,
            start_el=elements[0],
            open_el=elements[0],
            close_el=elements[-1],
        )
        return adapted, (lambda: None)

    host, children = _content_host(block, locate_mode)
    open_s = OxmlElement("w:p")
    close_s = OxmlElement("w:p")
    if children:
        children[0].addprevious(open_s)
        children[-1].addnext(close_s)
    else:
        # 空区间：宿主内插两个相邻哨兵
        host.insert(0, open_s)
        open_s.addnext(close_s)

    adapted = _AdaptedBlock(
        section_code=block.section_code,
        elements=[open_s, *children, close_s],
        start_el=open_s,
        open_el=open_s,
        close_el=close_s,
    )

    def _cleanup() -> None:
        # 持有引用删除（禁裸 id()）：哨兵只是插入 / 删除锚点，不进最终文档
        for s in (open_s, close_s):
            parent = s.getparent()
            if parent is not None:
                parent.remove(s)

    return adapted, _cleanup


def _has_rebuildable_structure(adapted: _AdaptedBlock, names: dict[str, str]) -> bool:
    """块内有**真实模板结构** ⇒ 按模板重建；否则纯文字回退（仍接过滤）。

    判据：开闭哨兵之间至少有一个 ``w:tbl``（需重渲染库内表），或有一个**标题样式**段落
    （``rebuild_block`` 会把它当块标题保留 / 重排编号）。

    为什么不认「有非空 ``w:p``」：交付件 / 测试里的纯文字块就是一段或几段普通正文，没有标题样式、
    没有表。``rebuild_block`` 的 ``_layout`` 在无标题样式时会把**首个普通文字段当标题**保留，
    导致刷新时旧首段文字留下、新正文另起 —— 旧正文没被换掉。这类块应走纯文字回退（整段按过滤后的
    DB 正文重渲染）。空区间（存量交付件锚点里什么都没有）同样回退。
    """
    from app.services.note_template_docx_base import _heading_level

    inner = adapted.elements[1:-1]
    for el in inner:
        if el.tag == _W_TBL:
            return True
        if el.tag == _W_P and _heading_level(el, names):
            return True
    return False


def rebuild_refresh_block(
    doc: Any,
    block: Any,
    note: Any,
    locate_mode: str,
    *,
    exporter: Any,
) -> RebuildOutcome:
    """按模板结构重建刷新块（标题 + 库内表 + 交付过滤后的正文）；无结构时纯文字回退（仍接过滤）。

    与全量生成（``fill_template_blocks``）同一套：``rebuild_block`` + ``make_block_filter``；正文过滤与
    结构判定共用同一组标题（块标题作 ``extra_titles``、块内子标题 / 表题作 ``structure_titles``）。

    Args:
        doc: python-docx Document。
        block: 定位块（``AnchorBlock`` 或 marker ``SectionBlock``）。
        note: 该章节 ``DisclosureNote`` ORM 行（调用方在 async 上下文已取好）。
        locate_mode: ``"anchor"`` / ``"marker"``。
        exporter: 提供 ``_render_table`` / ``_effective_table_data`` 的导出器（``NoteWordExporter``）。

    Returns:
        RebuildOutcome
    """
    import app.services.note_template_block_fill as tbf
    from app.services.note_deliverable_filter import FilterStats, make_block_filter
    from app.services.note_word_exporter import _template_body_run_style

    names = tbf.style_map(doc)
    styles = tbf.styles_by_id(doc)
    adapted, cleanup = _adapt(block, locate_mode)
    try:
        if not _has_rebuildable_structure(adapted, names):
            cleanup()
            return _fallback_text(doc, block, note, locate_mode)

        lay = tbf.block_layout(adapted, names)
        fstats = FilterStats()
        extra = (tbf.paragraph_text(lay.title),) if lay.title is not None else ()
        structure = tuple(tbf.structure_titles(lay))
        rebuilt = tbf.rebuild_block(
            doc,
            adapted,
            note,
            exporter=exporter,
            run_style=_template_body_run_style,
            text_filter=make_block_filter(
                note, fstats, extra_titles=extra, structure_titles=structure
            ),
            names=names,
            styles=styles,
            layout=lay,
        )
    finally:
        cleanup()

    return RebuildOutcome(
        rebuilt=True,
        tables_rendered=rebuilt.stats.tables_rendered,
        tables_cloned=rebuilt.stats.tables_cloned,
        dropped_lines=fstats.dropped_lines,
        dropped_tables=fstats.dropped_tables,
        placeholders_left=fstats.placeholders_left,
    )


def _fallback_text(
    doc: Any, block: Any, note: Any, locate_mode: str
) -> RebuildOutcome:
    """纯文字 / 无结构块：原位渲染过滤后的正文（接入交付过滤），返回统计。

    定位与原 ``_replace_block_content`` 一致（anchor 的内容控件宿主 / 空区间逆序插、marker 开标记后插），
    但正文经 ``make_block_filter`` 过滤（剔除大模型草稿 / 指引 / 占位），与程序化导出口径（Task 9）一致。
    """
    from app.services.note_deliverable_filter import FilterStats, make_block_filter
    from app.services.note_text_word_render import note_text_elements
    from app.services.section_anchor_utils import anchor_block_content_host

    text_content = (getattr(note, "text_content", None) or "")
    fstats = FilterStats()
    new_paragraphs = note_text_elements(
        text_content, block_filter=make_block_filter(note, fstats)
    )

    old_elements = list(getattr(block, "elements", None) or [])
    if locate_mode == "anchor":
        host, children = anchor_block_content_host(block)
        if children:
            anchor_el = children[0]
            for p in new_paragraphs:
                anchor_el.addprevious(p)
        else:
            start_el = block.start_el
            for p in reversed(new_paragraphs):
                start_el.addnext(p)
        for el in children:
            parent = el.getparent()
            if parent is not None:
                parent.remove(el)
    else:
        open_el = getattr(block, "open_el", None)
        inner = [
            el
            for el in old_elements
            if el is not open_el and el is not getattr(block, "close_el", None)
        ]
        if open_el is not None:
            for p in reversed(new_paragraphs):
                open_el.addnext(p)
        for el in inner:
            parent = el.getparent()
            if parent is not None:
                parent.remove(el)

    return RebuildOutcome(
        rebuilt=False,
        dropped_lines=fstats.dropped_lines,
        dropped_tables=fstats.dropped_tables,
        placeholders_left=fstats.placeholders_left,
    )

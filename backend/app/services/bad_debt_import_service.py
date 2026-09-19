"""坏账准备明细表 D2-3 Excel 导入服务（BadDebtImportService）

对应 design.md「Components and Interfaces #2 BadDebtImportService」。

两阶段导入：
1. parse_and_match: 解析 xlsx → 行匹配（与当前树 row_label 精确匹配）
2. commit_matched_rows: 批量更新匹配到的行金额（仅非 None 值覆盖）

口径铁律：
- service 只 flush 不 commit（router 统一 commit 保原子）。
- 金额列空值 = None（不覆盖原值），非空值覆盖原值。
- A 列去前导空格后精确匹配 row_label。

Requirements: 3.1, 3.2, 3.3, 3.6, 3.8, 5.3
"""

from __future__ import annotations

import uuid
from decimal import Decimal, InvalidOperation
from typing import Literal

from pydantic import BaseModel
from sqlalchemy.ext.asyncio import AsyncSession

from app.services.bad_debt_auto_sum import AutoSumEngine

# 13 金额列名 amount_b ~ amount_n
_AMOUNT_COLUMNS = AutoSumEngine.AMOUNT_COLUMNS


# ─── Pydantic 模型 ───────────────────────────────────────────────────────────


class ImportRowMatch(BaseModel):
    """单行匹配结果。"""

    excel_row_index: int  # Excel 行号（1-based）
    excel_label: str  # A 列原始项目名
    status: Literal["matched", "unmatched"]
    matched_row_id: uuid.UUID | None = None  # matched 时指向 bad_debt_detail_rows.id
    matched_row_label: str | None = None  # 匹配到的树行标签
    is_parent: bool = False  # 是否父行
    amounts: dict[str, Decimal | None] = {}  # B~N 列解析后的金额


class ImportParseResult(BaseModel):
    """解析+匹配结果。"""

    rows: list[ImportRowMatch]
    matched_count: int
    unmatched_count: int
    errors: list[str]  # 格式校验错误（非空时整体拒绝）


# ─── 服务类 ───────────────────────────────────────────────────────────────────


class BadDebtImportService:
    """坏账准备明细表 Excel 导入解析 + 匹配 + 批量写入。"""

    def __init__(self, db: AsyncSession):
        self.db = db

    async def parse_and_match(
        self, file_bytes: bytes, wp_index_id: uuid.UUID
    ) -> ImportParseResult:
        """解析 xlsx 文件并与当前坏账树 row_label 进行行匹配。

        1. 校验 xlsx 格式（>=14 列，A 列有数据）
        2. 跳过表头行（R1~R11），从 R12 开始读数据行
        3. A 列去前导空格后与当前树 row_label 精确匹配
        4. 解析 B~N 列金额（int/float/Decimal，空=None）

        Requirements: 3.1, 3.2, 3.3, 3.8
        """
        import io

        from openpyxl import load_workbook

        errors: list[str] = []

        # 解析 xlsx
        try:
            wb = load_workbook(
                filename=io.BytesIO(file_bytes), read_only=True, data_only=True
            )
        except Exception as exc:
            errors.append(f"无法解析 xlsx 文件: {exc}")
            return ImportParseResult(
                rows=[], matched_count=0, unmatched_count=0, errors=errors
            )

        ws = wb.active
        if ws is None:
            errors.append("xlsx 文件无活动工作表")
            wb.close()
            return ImportParseResult(
                rows=[], matched_count=0, unmatched_count=0, errors=errors
            )

        # 校验列数：至少 14 列（A~N）
        if ws.max_column is not None and ws.max_column < 14:
            errors.append(
                f"列数不足: 期望至少 14 列（A~N），实际 {ws.max_column} 列"
            )
            wb.close()
            return ImportParseResult(
                rows=[], matched_count=0, unmatched_count=0, errors=errors
            )

        # 获取当前树的 row_label 映射
        from app.services.bad_debt_nested_table_service import NestedTableService

        tree = await NestedTableService(self.db).get_tree(wp_index_id)
        label_map = self._build_label_map(tree)

        # 读取数据行（从 R12 开始，跳过 R1~R11 表头）
        rows: list[ImportRowMatch] = []
        data_start_row = 12

        for row_cells in ws.iter_rows(min_row=data_start_row):
            row_idx = row_cells[0].row  # type: ignore[union-attr]
            # A 列值
            a_value = row_cells[0].value  # type: ignore[union-attr]
            if a_value is None or str(a_value).strip() == "":
                continue  # 跳过空行

            excel_label = str(a_value)
            # 去前导空格后做匹配
            stripped_label = excel_label.lstrip()

            # 解析 B~N 列金额（列索引 1~13 对应 amount_b~amount_n）
            amounts: dict[str, Decimal | None] = {}
            for col_idx, col_name in enumerate(_AMOUNT_COLUMNS):
                cell_value = (
                    row_cells[col_idx + 1].value  # type: ignore[union-attr]
                    if col_idx + 1 < len(row_cells)
                    else None
                )
                amounts[col_name] = self._parse_cell_amount(cell_value)

            # 匹配
            match_info = label_map.get(stripped_label)
            if match_info is not None:
                row_match = ImportRowMatch(
                    excel_row_index=row_idx,
                    excel_label=excel_label,
                    status="matched",
                    matched_row_id=match_info["id"],
                    matched_row_label=match_info["label"],
                    is_parent=match_info["is_parent"],
                    amounts=amounts,
                )
            else:
                row_match = ImportRowMatch(
                    excel_row_index=row_idx,
                    excel_label=excel_label,
                    status="unmatched",
                    matched_row_id=None,
                    matched_row_label=None,
                    is_parent=False,
                    amounts=amounts,
                )
            rows.append(row_match)

        wb.close()

        # 校验 A 列必须有数据
        if not rows:
            errors.append("A 列无有效数据行（从第 12 行起）")

        matched_count = sum(1 for r in rows if r.status == "matched")
        unmatched_count = sum(1 for r in rows if r.status == "unmatched")

        return ImportParseResult(
            rows=rows,
            matched_count=matched_count,
            unmatched_count=unmatched_count,
            errors=errors,
        )

    async def commit_matched_rows(
        self, wp_index_id: uuid.UUID, rows: list[ImportRowMatch]
    ) -> int:
        """批量更新匹配到的行金额。

        - 仅处理 status="matched" 的行
        - 对每行，仅覆盖 amounts 中非 None 的金额列（空单元格不覆盖原值）
        - 返回实际更新行数

        Requirements: 3.6, 5.3
        """
        from app.services.bad_debt_nested_table_service import NestedTableService

        svc = NestedTableService(self.db)
        updated_count = 0

        for row in rows:
            if row.status != "matched":
                continue
            if row.matched_row_id is None:
                continue

            # 获取 ORM 行
            db_row = await svc._get_row(row.matched_row_id)

            # 仅覆盖非 None 的金额列
            has_update = False
            for col_name in _AMOUNT_COLUMNS:
                new_value = row.amounts.get(col_name)
                if new_value is not None:
                    setattr(db_row, col_name, new_value)
                    has_update = True

            if has_update:
                db_row.version += 1
                updated_count += 1

                # 子行变更 → 重算其父行
                if db_row.parent_row_id is not None:
                    parent = await svc._get_row(db_row.parent_row_id)
                    await svc._resum_parent(parent)

        await self.db.flush()
        return updated_count

    # ─── 内部工具 ────────────────────────────────────────────────────────────

    @staticmethod
    def _build_label_map(tree) -> dict[str, dict]:
        """从树结构构建 row_label → {id, label, is_parent} 映射。

        父行和子行都参与匹配。标签去重时后出现的覆盖先出现的（不太可能重复）。
        """
        label_map: dict[str, dict] = {}

        for parent in tree.parents:
            label_map[parent.row_label] = {
                "id": parent.id,
                "label": parent.row_label,
                "is_parent": True,
            }
            for child in parent.children:
                label_map[child.row_label] = {
                    "id": child.id,
                    "label": child.row_label,
                    "is_parent": False,
                }

        return label_map

    @staticmethod
    def _parse_cell_amount(value) -> Decimal | None:
        """解析单元格金额值为 Decimal，空值返回 None。

        支持 int/float/str 格式的数值。
        """
        if value is None:
            return None
        if isinstance(value, str):
            value = value.strip()
            if value == "":
                return None
        try:
            dec = Decimal(str(value))
            if dec.is_nan() or dec.is_infinite():
                return None
            return dec
        except (InvalidOperation, ValueError, TypeError):
            return None

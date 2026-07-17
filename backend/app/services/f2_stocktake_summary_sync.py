"""F2-23 存货监盘小结 — 结构化字段 ↔ G2-6-1 风格 docx 双向同步.

结构化真源：checklist_responses.item_id = ``F2-23-fields``（remark=JSON）。
在线编辑真源：项目 OnlyOffice 缓存 ``{storage}/F2-23.docx``。
"""

from __future__ import annotations

import json
import logging
import re
import shutil
from pathlib import Path
from typing import Any

logger = logging.getLogger(__name__)

FIELDS_ITEM_ID = "F2-23-fields"
NOTE_ITEM_ID = "F2-23-note"

F2_23_SYNC_FIELDS: tuple[str, ...] = (
    "entityName",
    "auditYear",
    "bsDate",
    "purpose",
    "scope",
    "warehouses",
    "countDate",
    "clientStaff",
    "auditors",
    "assignment",
    "countMethod",
    "processOverview",
    "inventoryTotal",
    "sampleAmount",
    "samplePct",
    "coverageNote",
    "resultByLocation",
    "overallConclusion",
    "followUp",
    "teamName",
    "summaryDate",
)

_PLACEHOLDER_RE = re.compile(r"\$\{([a-zA-Z_][a-zA-Z0-9_]*)\}")

_SECTION_EXTRACTORS: list[tuple[str, re.Pattern[str]]] = [
    (
        "purpose",
        re.compile(
            r"1[．.\s]*监盘目的[：:]\s*(.*?)(?=\n\s*2[．.\s]*监盘范围|\Z)",
            re.S,
        ),
    ),
    (
        "scope",
        re.compile(
            r"2[．.\s]*监盘范围[：:]\s*(.*?)(?=\n\s*3[．.\s]*监盘地点|\Z)",
            re.S,
        ),
    ),
    (
        "warehouses",
        re.compile(
            r"3[．.\s]*监盘地点[：:]\s*(.*?)(?=\n\s*4[．.\s]*监盘时间|\Z)",
            re.S,
        ),
    ),
    (
        "countDate",
        re.compile(
            r"4[．.\s]*监盘时间[：:]\s*(.*?)(?=\n\s*5[．.\s]*参与人员|\Z)",
            re.S,
        ),
    ),
    (
        "countMethod",
        re.compile(
            r"6[．.\s]*公司存货盘点方法[：:]\s*(.*?)(?=\n\s*7[．.\s]*监盘情况汇总|\Z)",
            re.S,
        ),
    ),
    (
        "processOverview",
        re.compile(
            r"7[．.\s]*监盘情况汇总[：:]\s*(.*?)(?=\n\s*8[．.\s]*具体监盘情况|\Z)",
            re.S,
        ),
    ),
    (
        "resultByLocation",
        re.compile(
            r"8[．.\s]*具体监盘情况[：:]\s*(.*?)(?=\n\s*9[．.\s]*监盘结论|\Z)",
            re.S,
        ),
    ),
    (
        "overallConclusion",
        re.compile(
            r"9[．.\s]*监盘结论[：:]\s*(.*?)(?=\n\s*附[：:]|\n\s*致同|\Z)",
            re.S,
        ),
    ),
]


def _format_cn_date(val: str) -> str:
    s = (val or "").strip()
    if not s:
        return s
    m = re.match(r"^(\d{4})-(\d{1,2})-(\d{1,2})$", s)
    if m:
        return f"{m.group(1)}年{int(m.group(2))}月{int(m.group(3))}日"
    return s


def _normalize_iso_date(val: str) -> str:
    s = (val or "").strip()
    if not s or "[" in s:
        return ""
    if re.match(r"^\d{4}-\d{2}-\d{2}$", s):
        return s
    m = re.search(r"(\d{4})\s*年\s*(\d{1,2})\s*月\s*(\d{1,2})\s*日?", s)
    if m:
        return f"{m.group(1)}-{int(m.group(2)):02d}-{int(m.group(3)):02d}"
    m2 = re.match(r"^(\d{4})[./-](\d{1,2})[./-](\d{1,2})", s)
    if m2:
        return f"{m2.group(1)}-{int(m2.group(2)):02d}-{int(m2.group(3)):02d}"
    return s


def parse_fields_json(raw: str | None) -> dict[str, str]:
    if not raw:
        return {}
    try:
        data = json.loads(raw)
    except (TypeError, json.JSONDecodeError):
        return {}
    if not isinstance(data, dict):
        return {}
    return {str(k): ("" if v is None else str(v)) for k, v in data.items()}


def migrate_legacy_fields(fields: dict[str, str]) -> dict[str, str]:
    """旧 F2-23 键迁移到 G2-6-1 字段。"""
    out = dict(fields)
    pairs = (
        ("inventoryComposition", "countMethod"),
        ("sampleCoverage", "coverageNote"),
        ("opinionImpact", "followUp"),
    )
    for src, dst in pairs:
        if (out.get(src) or "").strip() and not (out.get(dst) or "").strip():
            out[dst] = out[src]
    return out


def build_placeholder_map(
    fields: dict[str, str],
    *,
    project_context: dict[str, Any] | None = None,
) -> dict[str, str]:
    ctx = project_context or {}
    merged = migrate_legacy_fields(dict(fields))
    if not merged.get("entityName") and ctx.get("client_name"):
        merged["entityName"] = str(ctx["client_name"])
    if not merged.get("auditYear") and ctx.get("audit_year"):
        merged["auditYear"] = str(ctx["audit_year"])
    if not merged.get("bsDate") and ctx.get("bs_date"):
        merged["bsDate"] = str(ctx["bs_date"])
    if not merged.get("teamName"):
        merged["teamName"] = "XXXX审计小组"

    # 第七节：过程 + 覆盖数字 + 说明合并为可阅读段落（docx 单占位）
    overview = (merged.get("processOverview") or "").strip()
    total = (merged.get("inventoryTotal") or "").strip()
    sample = (merged.get("sampleAmount") or "").strip()
    pct = (merged.get("samplePct") or "").strip()
    note = (merged.get("coverageNote") or "").strip()
    coverage_bits: list[str] = []
    if total or sample or pct:
        coverage_bits.append(
            f"存货总额约{total or '____'}万元，抽盘金额约{sample or '____'}万元，"
            f"占比约{pct or '____'}%。"
        )
    if note:
        coverage_bits.append(note)
    summary_block = overview
    if coverage_bits:
        summary_block = (
            f"{overview}\n{' '.join(coverage_bits)}".strip()
            if overview
            else " ".join(coverage_bits)
        )
    if summary_block:
        merged["summaryBlock"] = summary_block

    conclusion = (merged.get("overallConclusion") or "").strip()
    follow = (merged.get("followUp") or "").strip()
    if follow:
        merged["conclusionBlock"] = (
            f"{conclusion}\n需进一步跟进：{follow}".strip()
            if conclusion
            else f"需进一步跟进：{follow}"
        )
    else:
        merged["conclusionBlock"] = conclusion

    result: dict[str, str] = {}
    for fid in (*F2_23_SYNC_FIELDS, "summaryBlock", "conclusionBlock"):
        token = f"${{{fid}}}"
        val = (merged.get(fid) or "").strip()
        if fid in ("summaryDate", "bsDate") and val:
            val = _format_cn_date(val)
        result[token] = val if val else token
    return result


def fill_summary_docx(
    template_path: Path,
    target_path: Path,
    fields: dict[str, str],
    *,
    project_context: dict[str, Any] | None = None,
) -> Path:
    from docx import Document

    target_path.parent.mkdir(parents=True, exist_ok=True)
    shutil.copy2(template_path, target_path)
    doc = Document(str(target_path))
    replacements = build_placeholder_map(fields, project_context=project_context)

    def _replace_paragraph_text(paragraph: Any) -> None:
        text = paragraph.text or ""
        if not text or "${" not in text:
            return
        new_text = text
        for key, val in replacements.items():
            if key in new_text:
                new_text = new_text.replace(key, val)
        if new_text == text:
            return
        for run in paragraph.runs:
            run.text = ""
        if paragraph.runs:
            paragraph.runs[0].text = new_text
        else:
            paragraph.add_run(new_text)

    for paragraph in doc.paragraphs:
        _replace_paragraph_text(paragraph)
    for table in doc.tables:
        for row in table.rows:
            for cell in row.cells:
                for paragraph in cell.paragraphs:
                    _replace_paragraph_text(paragraph)

    doc.save(str(target_path))
    return target_path


def _clean_section_body(text: str) -> str:
    t = (text or "").strip()
    t = _PLACEHOLDER_RE.sub("", t).strip()
    return t


def extract_fields_from_docx(docx_path: Path) -> dict[str, str]:
    from docx import Document

    doc = Document(str(docx_path))
    paras = [p.text.strip() for p in doc.paragraphs if p.text is not None]
    nonempty = [t for t in paras if t]
    full = "\n".join(nonempty)
    out: dict[str, str] = {}

    for i, t in enumerate(nonempty):
        if t.startswith("索引号"):
            if i + 1 < len(nonempty) and "监盘小结" not in nonempty[i + 1]:
                cand = nonempty[i + 1]
                if not cand.startswith("根据") and "${" not in cand:
                    out["entityName"] = cand
            if i + 2 < len(nonempty):
                m_year = re.search(r"(\d{4})\s*年存货监盘小结", nonempty[i + 2])
                if m_year:
                    out["auditYear"] = m_year.group(1)
            # 也可能标题同行
            m_title = re.search(r"(\d{4})\s*年存货监盘小结", t)
            if m_title:
                out["auditYear"] = m_title.group(1)
            break
    # 标题行可能是「单位名」下一行
    for t in nonempty[:5]:
        m_year = re.search(r"(\d{4})\s*年存货监盘小结", t)
        if m_year:
            out["auditYear"] = m_year.group(1)
            break

    m_bs = re.search(r"(\d{4}[-年]\d{1,2}[-月]\d{1,2}日?)\s*资产负债表", full)
    if m_bs:
        out["bsDate"] = _normalize_iso_date(
            m_bs.group(1).replace("年", "-").replace("月", "-").replace("日", "")
        ) or _normalize_iso_date(m_bs.group(1))

    for field_id, pattern in _SECTION_EXTRACTORS:
        m = pattern.search(full)
        if not m:
            continue
        body = _clean_section_body(m.group(1))
        if not body or body.startswith("${"):
            continue
        if field_id == "processOverview":
            # 尝试拆出金额覆盖
            m_amt = re.search(
                r"存货总额[约]?(\d[\d.,]*)\s*万元.*?抽盘金额[约]?(\d[\d.,]*)\s*万元.*?占比[约]?(\d[\d.,]*)\s*%",
                body,
                re.S,
            )
            if m_amt:
                out["inventoryTotal"] = m_amt.group(1)
                out["sampleAmount"] = m_amt.group(2)
                out["samplePct"] = m_amt.group(3)
            # 覆盖说明：括号/「因…」句
            m_note = re.search(r"[（(]([^）)]*未结账[^）)]*)[）)]", body)
            if m_note:
                out["coverageNote"] = m_note.group(1).strip()
            out["processOverview"] = body
        elif field_id == "overallConclusion":
            m_follow = re.search(r"需进一步跟进[：:]\s*(.*)", body, re.S)
            if m_follow:
                out["followUp"] = _clean_section_body(m_follow.group(1))
                out["overallConclusion"] = _clean_section_body(body[: m_follow.start()])
            else:
                out["overallConclusion"] = body
        else:
            out[field_id] = body

    # 第 5 节人员
    m5 = re.search(
        r"5[．.\s]*参与人员及分工[：:]?\s*(.*?)(?=\n\s*6[．.\s]*公司存货盘点方法|\Z)",
        full,
        re.S,
    )
    if m5:
        block = m5.group(1)
        m_cli = re.search(r"(?:盘点人员|被审计单位)[^\n：:]*[：:]\s*(.+)", block)
        m_aud = re.search(r"项目组监盘人员[：:]\s*(.+)", block)
        m_asg = re.search(r"分工[：:]\s*(.*)", block, re.S)
        if m_cli:
            out["clientStaff"] = m_cli.group(1).strip()
        if m_aud:
            out["auditors"] = m_aud.group(1).strip()
        if m_asg:
            out["assignment"] = _clean_section_body(m_asg.group(1))
        elif not out.get("assignment"):
            # 无「分工：」标题时整段作为分工
            cleaned = _clean_section_body(block)
            if cleaned and not cleaned.startswith("${"):
                out["assignment"] = cleaned

    m_sign = re.search(
        r"致同会计师事务所\n([^\n]+)\n([^\n]*20\d{2}\s*年\s*\d{1,2}\s*月\s*\d{1,2}\s*日)",
        full,
    )
    if m_sign:
        team = m_sign.group(1).strip()
        if team and not team.startswith("${") and "[" not in team:
            out["teamName"] = team
        plan_raw = re.sub(r"\s+", "", m_sign.group(2))
        plan_iso = _normalize_iso_date(plan_raw)
        if plan_iso:
            out["summaryDate"] = plan_iso

    for fid in list(out.keys()):
        if not out[fid] or out[fid].startswith("${") or out[fid] == f"${{{fid}}}":
            del out[fid]

    return out


def merge_extracted_into_existing(
    existing: dict[str, str],
    extracted: dict[str, str],
) -> dict[str, str]:
    merged = migrate_legacy_fields(dict(existing))
    for k, v in extracted.items():
        if v and v.strip():
            merged[k] = v.strip()
    return merged


def create_g2_6_1_template_docx(target: Path) -> Path:
    """生成带 ``${}`` 占位符的 F2-23 / G2-6-1 结构模板。"""
    from docx import Document
    from docx.oxml.ns import qn
    from docx.shared import Pt, RGBColor

    doc = Document()
    style = doc.styles["Normal"]
    style.font.name = "宋体"
    style.font.size = Pt(12)
    style._element.rPr.rFonts.set(qn("w:eastAsia"), "宋体")

    def add(text: str, *, bold: bool = False, red: bool = False) -> None:
        p = doc.add_paragraph()
        run = p.add_run(text)
        run.bold = bold
        if red:
            run.font.color.rgb = RGBColor(0xC0, 0x00, 0x00)

    add("索引号：G2-6-1")
    add("${entityName}", bold=True, red=True)
    add("${auditYear}年存货监盘小结", bold=True)
    add("")
    add(
        "根据《中国注册会计师审计准则第1311号——存货监盘》的要求，"
        "为了对${entityName}（以下简称该公司或公司）${bsDate}资产负债表上存货的"
        "存在性、所有权、品质状况等情况进行验证，"
        "项目组对该公司存货实施了监盘，现将监盘情况小结如下。"
    )
    add("1．监盘目的：", bold=True)
    add("${purpose}")
    add("2．监盘范围：", bold=True)
    add("${scope}")
    add("3．监盘地点：", bold=True)
    add("${warehouses}")
    add("4．监盘时间：", bold=True)
    add("${countDate}")
    add("5．参与人员及分工：", bold=True)
    add("盘点人员：${clientStaff}")
    add("项目组监盘人员：${auditors}")
    add("分工：")
    add("${assignment}")
    add("6．公司存货盘点方法：", bold=True)
    add("${countMethod}")
    add("7．监盘情况汇总：", bold=True)
    add("${summaryBlock}")
    add("8．具体监盘情况：", bold=True)
    add("${resultByLocation}")
    add("9．监盘结论：", bold=True)
    add("${conclusionBlock}")
    add("")
    add("附：各分厂抽盘明细表")
    add("")
    add("致同会计师事务所")
    add("${teamName}")
    add("${summaryDate}")

    target.parent.mkdir(parents=True, exist_ok=True)
    doc.save(str(target))
    return target

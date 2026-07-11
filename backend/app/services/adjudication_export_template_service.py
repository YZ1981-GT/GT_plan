"""审定表导出模板动态生成器。

动态生成含多行合并表头、编制说明 sheet、行骨架的空白 xlsx 模板，
替代现有 export-template 端点直接返回磁盘源 xlsx 的行为。
全部 D~N 审定表（``^[A-N]\\d+-1$``）复用同一生成器。
"""

from __future__ import annotations

import io
import logging
import os
import re
from typing import TYPE_CHECKING

from openpyxl.styles import Alignment, Border, Font, Side

if TYPE_CHECKING:
    from sqlalchemy.ext.asyncio import AsyncSession

logger = logging.getLogger(__name__)

# 审定表 wp_code 匹配正则（编译一次复用）
_ADJUDICATION_RE = re.compile(r"^[A-N]\d+-1$")


class AdjudicationExportTemplateService:
    """审定表导出模板动态生成器服务。

    纯函数 + async 账龄解析，只读无写库。
    """

    # 含账龄列的 wp_code 前缀（去掉 -1 后缀后的科目标识）
    AGING_SUBJECTS: dict[str, str] = {
        "D2": "D2",
        "D3": "D3",
        "F1": "F1",
        "K1": "K1",
        "K3": "K3",
        "G5": "G5",
    }

    @staticmethod
    def is_adjudication_table(wp_code: str) -> bool:
        """判断 wp_code 是否为审定表（匹配 ^[A-N]\\d+-1$）。"""
        return bool(_ADJUDICATION_RE.match(wp_code or ""))

    @staticmethod
    def _get_aging_subject(wp_code: str) -> str | None:
        """从 wp_code 提取账龄科目标识。

        例如 "D2-1" → "D2"，若前缀在 AGING_SUBJECTS 中则返回，否则返回 None。
        """
        prefix = (wp_code or "").replace("-1", "")
        return prefix if prefix in AdjudicationExportTemplateService.AGING_SUBJECTS else None

    @staticmethod
    async def _resolve_aging_headers(db: "AsyncSession", wp_id: str, subject: str) -> list[str]:
        """为账龄科目解析动态列头。

        集成 resolve_aging_segments + subject_aging_periods + build_aging_headers。
        失败时回退到默认预设段（保证不返回空列表）。
        """
        from app.routers.wp_render_strategies._cycle_import_export_common import (
            build_aging_headers,
            resolve_aging_segments,
            subject_aging_periods,
        )

        try:
            segments = await resolve_aging_segments(db, wp_id, subject)
            periods = subject_aging_periods(subject)
            if not segments:
                raise ValueError(f"No segments resolved for subject={subject}")
            return build_aging_headers(segments, periods)
        except Exception as e:
            logger.warning(
                "_resolve_aging_headers: 账龄段解析失败, subject=%s, wp_id=%s: %s. 使用默认预设。",
                subject,
                wp_id,
                e,
            )
            # Double fallback: use default preset directly
            from app.services import aging_config_service as _acs

            default_preset = _acs.DEFAULT_SUBJECT_PRESETS.get(subject, _acs.AgingPreset.FIVE_YEAR)
            segments = _acs.resolve_segments(default_preset, None)
            periods = subject_aging_periods(subject)
            return build_aging_headers(segments, periods)

    @staticmethod
    def _build_instruction_sheet(ws, wp_code: str, wp_name: str, has_aging: bool) -> None:
        """填充编制说明 sheet 内容。"""
        bold = Font(bold=True)
        row = 1

        # ── 标题 ──
        ws.cell(row=row, column=1, value=f"{wp_name}审定表 ({wp_code}) — 编制说明").font = bold
        row += 2

        # ── 一、各列含义 (Sub-task 4.1) ──
        ws.cell(row=row, column=1, value="一、各列含义").font = bold
        row += 1

        # 表头
        for ci, h in enumerate(["列号", "列名", "说明"], 1):
            ws.cell(row=row, column=ci, value=h).font = bold
        row += 1

        columns_info = [
            (1, "项目", "科目行名称（预填）"),
            (2, "期初未审", "期初未审数（手填）"),
            (3, "期初AJE", "期初审计调整（手填）"),
            (4, "期初RJE", "期初重分类调整（手填）"),
            (5, "期初审定", "期初审定数（自动计算=未审+AJE+RJE）"),
            (6, "期末未审", "期末未审数（手填）"),
            (7, "期末AJE", "期末审计调整（手填）"),
            (8, "期末RJE", "期末重分类调整（手填）"),
            (9, "期末审定", "期末审定数（自动计算=未审+AJE+RJE）"),
            (10, "变动额", "变动金额（自动计算=期末审定-期初审定）"),
            (11, "变动率", "变动比率（自动计算=变动额/期初审定）"),
            (12, "原因分析", "重大变动原因说明（手填）"),
        ]
        for col_no, col_name, col_desc in columns_info:
            ws.cell(row=row, column=1, value=col_no)
            ws.cell(row=row, column=2, value=col_name)
            ws.cell(row=row, column=3, value=col_desc)
            row += 1

        # 账龄注释
        if has_aging:
            row += 1
            ws.cell(
                row=row,
                column=1,
                value='注：含账龄列的审定表在\u201c期末审定\u201d与\u201c变动额\u201d之间会有动态账龄列，列数取决于项目账龄配置。',
            )
            row += 1

        row += 1

        # ── 二、只读自动计算列 (Sub-task 4.2) ──
        ws.cell(row=row, column=1, value="二、只读自动计算列").font = bold
        row += 1

        formulas = [
            "期初审定 = 期初未审 + 期初AJE + 期初RJE",
            "期末审定 = 期末未审 + 期末AJE + 期末RJE",
            "变动额 = 期末审定 - 期初审定",
            "变动率 = 变动额 / 期初审定",
        ]
        for f in formulas:
            ws.cell(row=row, column=1, value=f"• {f}")
            row += 1

        ws.cell(
            row=row,
            column=1,
            value="以上列由系统自动计算，导入时即使模板中有值也不导入。",
        )
        row += 2

        # ── 三、填报规则 (Sub-task 4.3) ──
        ws.cell(row=row, column=1, value="三、填报规则").font = bold
        row += 1

        ws.cell(
            row=row,
            column=1,
            value="用户需填写列：期初未审、期初AJE、期初RJE、期末未审、期末AJE、期末RJE、原因分析",
        )
        row += 1
        ws.cell(
            row=row,
            column=1,
            value="只读自动计算列（导入时忽略）：期初审定、期末审定、变动额、变动率",
        )
        row += 2

        # ── 四、导入注意事项 (Sub-task 4.4) ──
        ws.cell(row=row, column=1, value="四、导入注意事项").font = bold
        row += 1

        import_notes = [
            '行匹配规则：按\u201c项目\u201d列名称精确匹配',
            "只读列：导入时忽略（即使模板中有值也不导入）",
            "空行：自动跳过",
            "顺序：不要求行顺序与模板一致",
        ]
        for note in import_notes:
            ws.cell(row=row, column=1, value=f"• {note}")
            row += 1

        # ── 设置列宽提升可读性 ──
        ws.column_dimensions["A"].width = 8
        ws.column_dimensions["B"].width = 15
        ws.column_dimensions["C"].width = 60

    @staticmethod
    def _build_data_sheet(
        ws,
        wp_code: str,
        wp_name: str,
        row_skeleton: list[dict],
        aging_headers: list[str] | None = None,
    ) -> None:
        """填充数据模板 sheet（标题+表头+行骨架）。"""
        # ── Row 1: Title ──
        title = f"{wp_name}审定表 {wp_code}"
        ws.cell(row=1, column=1, value=title).font = Font(bold=True, size=14)
        # Merge title across all columns (12 base + aging)
        total_cols = 12 + len(aging_headers or [])
        ws.merge_cells(start_row=1, start_column=1, end_row=1, end_column=total_cols)

        # ── Rows 2-3: Multi-row header ──
        AdjudicationExportTemplateService._build_multi_row_header(
            ws, start_col=1, aging_headers=aging_headers
        )

        # ── Rows 4+: Row skeleton (project column) ──
        data_start_row = 4
        bold_font = Font(bold=True)
        for i, item in enumerate(row_skeleton):
            row_num = data_start_row + i
            cell = ws.cell(row=row_num, column=1, value=item.get("item", ""))
            if item.get("is_section") or item.get("is_total"):
                cell.font = bold_font

    @staticmethod
    def _build_multi_row_header(
        ws,
        start_col: int,
        aging_headers: list[str] | None = None,
    ) -> None:
        """构建 2-3 行多行合并表头。

        Row 2: 一级表头（项目/期初/期末/[账龄列]/变动额/变动率/原因分析）
        Row 3: 二级表头（期初/期末各下辖 未审/AJE/RJE/审定）

        合并规则：
        - 项目/变动额/变动率/原因分析: 纵向合并 rows 2-3
        - 期初/期末: 横向合并各 4 列 (row 2)
        - 账龄列: 纵向合并 rows 2-3（每列独立，无父级分组）
        """
        thin_border = Border(
            left=Side(style="thin"),
            right=Side(style="thin"),
            top=Side(style="thin"),
            bottom=Side(style="thin"),
        )
        header_font = Font(bold=True)
        center_align = Alignment(horizontal="center", vertical="center")

        aging_cols = aging_headers or []

        # ── Sub-task 2.1: Row 2 level-1 headers ──────────────────────────
        # 固定列: 项目(1) + 期初(2-5) + 期末(6-9) = 9 列
        aging_start_col = 10  # 账龄列起始位置（紧接期末审定之后）
        tail_start = aging_start_col + len(aging_cols)  # 变动额起始列

        ws.cell(row=2, column=1, value="项目")
        ws.cell(row=2, column=2, value="期初")
        ws.cell(row=2, column=6, value="期末")

        # 账龄列写入 row 2（每列独立，纵向合并 rows 2-3）
        for i, ah in enumerate(aging_cols):
            ws.cell(row=2, column=aging_start_col + i, value=ah)

        ws.cell(row=2, column=tail_start, value="变动额")
        ws.cell(row=2, column=tail_start + 1, value="变动率")
        ws.cell(row=2, column=tail_start + 2, value="原因分析")

        # ── Sub-task 2.2: Row 3 level-2 headers ──────────────────────────
        sub_headers = ["未审", "AJE", "RJE", "审定"]
        for i, h in enumerate(sub_headers):
            ws.cell(row=3, column=2 + i, value=h)  # 期初: cols 2-5
            ws.cell(row=3, column=6 + i, value=h)  # 期末: cols 6-9

        # ── Sub-task 2.3: Merge operations ────────────────────────────────
        # 横向合并 (row 2)
        ws.merge_cells(start_row=2, start_column=2, end_row=2, end_column=5)  # 期初
        ws.merge_cells(start_row=2, start_column=6, end_row=2, end_column=9)  # 期末

        # 纵向合并 (rows 2-3)
        ws.merge_cells(start_row=2, start_column=1, end_row=3, end_column=1)  # 项目
        ws.merge_cells(
            start_row=2, start_column=tail_start, end_row=3, end_column=tail_start
        )  # 变动额
        ws.merge_cells(
            start_row=2, start_column=tail_start + 1, end_row=3, end_column=tail_start + 1
        )  # 变动率
        ws.merge_cells(
            start_row=2, start_column=tail_start + 2, end_row=3, end_column=tail_start + 2
        )  # 原因分析

        # 账龄列纵向合并 (rows 2-3, 每列独立)
        for i in range(len(aging_cols)):
            col = aging_start_col + i
            ws.merge_cells(start_row=2, start_column=col, end_row=3, end_column=col)

        # ── Sub-task 2.4: Apply styles ────────────────────────────────────
        total_cols = tail_start + 2  # 最后一列 (原因分析)
        for row in (2, 3):
            for col in range(1, total_cols + 1):
                cell = ws.cell(row=row, column=col)
                cell.font = header_font
                cell.alignment = center_align
                cell.border = thin_border

    @staticmethod
    def _get_row_skeleton(template_file_path: str | None, wp_code: str) -> list[dict]:
        """从模板 xlsx 提取审定表行骨架（降级返回 []）。

        1. template_file_path 为 None 或文件不存在 → 返回 []
        2. 构造 sheet_name，尝试 "审定表{wp_code}" 优先，回退到含 wp_code 的第一个 sheet
        3. 调用 extract_audit_rows 提取行骨架
        4. 任何异常 → 返回 []（降级）
        """
        if not template_file_path or not os.path.isfile(template_file_path):
            logger.debug(
                "_get_row_skeleton: 模板文件不存在或未提供, path=%s", template_file_path
            )
            return []

        try:
            from app.services.wp_audit_sheet_extract import extract_audit_rows
            from app.services.xlsx_read_adapter import list_sheet_names

            # 标准命名: "审定表{wp_code}"（如 "审定表D2-1"）
            sheet_name = f"审定表{wp_code}"

            # 检查 sheet 是否存在，若不存在则回退到含 wp_code 的第一个 sheet
            names = list_sheet_names(template_file_path)
            if sheet_name not in names:
                # 回退: 查找含 wp_code 的 sheet（兼容 "长期股权投资审定表G7-1" 等变体）
                # 优先匹配以 wp_code 结尾的 sheet（精确），其次包含 wp_code 的 sheet
                candidates = [n for n in names if n.rstrip().endswith(wp_code)]
                if not candidates:
                    candidates = [n for n in names if wp_code in n]
                if not candidates:
                    logger.warning(
                        "_get_row_skeleton: 未找到匹配 sheet, wp_code=%s, sheets=%s",
                        wp_code,
                        names[:10],
                    )
                    return []
                sheet_name = candidates[0]

            rows = extract_audit_rows(template_file_path, sheet_name)
            logger.debug(
                "_get_row_skeleton: 提取行骨架成功, wp_code=%s, sheet=%s, rows=%d",
                wp_code,
                sheet_name,
                len(rows),
            )
            return rows

        except Exception as e:
            logger.warning(
                "_get_row_skeleton: 提取行骨架失败, wp_code=%s, path=%s: %s",
                wp_code,
                template_file_path,
                e,
            )
            return []

    @staticmethod
    async def generate(
        wp_id: str,
        wp_code: str,
        wp_name: str,
        db: "AsyncSession",
        template_file_path: str | None = None,
    ) -> tuple[io.BytesIO, str]:
        """生成审定表 xlsx 模板。

        Returns:
            (xlsx_bytes_io, filename)
        """
        from openpyxl import Workbook

        cls = AdjudicationExportTemplateService

        # 1. Aging resolution
        aging_subject = cls._get_aging_subject(wp_code)
        aging_headers: list[str] | None = None
        if aging_subject:
            aging_headers = await cls._resolve_aging_headers(db, wp_id, aging_subject)

        # 2. Row skeleton
        row_skeleton = cls._get_row_skeleton(template_file_path, wp_code)

        # 3. Create workbook
        wb = Workbook()

        # 4. Instruction sheet (first sheet, rename default)
        ws_instruction = wb.active
        ws_instruction.title = "编制说明"
        cls._build_instruction_sheet(ws_instruction, wp_code, wp_name, has_aging=bool(aging_subject))

        # 5. Data sheet (second sheet)
        ws_data = wb.create_sheet(title=f"{wp_name}审定表")
        cls._build_data_sheet(ws_data, wp_code, wp_name, row_skeleton, aging_headers)

        # 6. Save to BytesIO
        buf = io.BytesIO()
        wb.save(buf)
        buf.seek(0)

        # 7. Filename
        filename = f"{wp_code}_{wp_name}审定表_模板.xlsx"

        return buf, filename

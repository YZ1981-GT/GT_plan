"""a18_summary_generator — A18-1 审计小结生成

从 A17-1 章节内容 + KAM 生成审计小结框架。
依赖 A17-core（A17-1 章节数据 + A17-2-1 KAM 可选）。

设计（见 a18 requirements）：
- 读取 A17-1 关键章节
- 附加 A17-2-1 KAM 摘要（如有）
- 返回结构化 sections 供弹窗预览 / prefilled-download 写入 docx
"""

from __future__ import annotations

import json
import logging
from typing import TYPE_CHECKING
from uuid import UUID

import sqlalchemy as sa
from sqlalchemy.ext.asyncio import AsyncSession

from app.models.workpaper_models import WorkingPaper, WpIndex

if TYPE_CHECKING:
    from docx.document import Document

logger = logging.getLogger(__name__)

_KEY_CHAPTERS = [
    ("A17-1-ch01", "审计业务约定范围"),
    ("A17-1-ch06", "重大错报风险应对"),
    ("A17-1-ch08", "已审报表分析"),
    ("A17-1-ch12", "关键审计事项"),
    ("A17-1-ch14", "审计结论"),
]


async def _find_wp_id(db: AsyncSession, project_id: UUID, wp_code: str) -> UUID | None:
    stmt = (
        sa.select(WorkingPaper.id)
        .join(WpIndex, WorkingPaper.wp_index_id == WpIndex.id)
        .where(
            WorkingPaper.project_id == project_id,
            WorkingPaper.is_deleted == sa.false(),
            WpIndex.wp_code == wp_code,
        )
        .limit(1)
    )
    return (await db.execute(stmt)).scalar_one_or_none()


async def _load_chapter_remarks(
    db: AsyncSession, wp_id: UUID, chapter_ids: list[str]
) -> dict[str, str]:
    if not chapter_ids:
        return {}
    stmt = sa.text(
        "SELECT item_id, remark FROM checklist_responses "
        "WHERE wp_id = :wp_id AND item_id = ANY(:ids)"
    )
    result = await db.execute(
        stmt,
        {"wp_id": str(wp_id), "ids": chapter_ids},
    )
    return {row.item_id: (row.remark or "") for row in result.mappings().all()}


async def _load_kam_summary(db: AsyncSession, project_id: UUID) -> str:
    kam_wp_id = await _find_wp_id(db, project_id, "A17-2-1")
    if not kam_wp_id:
        return ""

    stmt = sa.text(
        "SELECT item_id, conclusion, remark FROM checklist_responses "
        "WHERE wp_id = :wp_id AND item_id LIKE 'A17-2-1-KAM-%' "
        "ORDER BY item_id"
    )
    rows = (await db.execute(stmt, {"wp_id": str(kam_wp_id)})).mappings().all()
    if not rows:
        return ""

    lines: list[str] = []
    for row in rows:
        title = (row["conclusion"] or "").strip()
        if not title:
            continue
        lines.append(f"- {title}")
        remark_raw = row["remark"]
        if remark_raw:
            try:
                remark = json.loads(remark_raw)
                for key, label in (
                    ("situation", "事项"),
                    ("reason", "确定为KAM的原因"),
                    ("response", "审计应对"),
                ):
                    val = (remark.get(key) or "").strip()
                    if val:
                        lines.append(f"  {label}：{val}")
            except (json.JSONDecodeError, TypeError):
                pass
    return "\n".join(lines)


async def generate_audit_summary(db: AsyncSession, project_id: UUID) -> dict:
    """从 A17-1 章节 + KAM 生成审计小结框架。"""
    wp_id = await _find_wp_id(db, project_id, "A17-1")
    if not wp_id:
        logger.info("A18-1 生成：未找到 A17-1 底稿 (project=%s)", project_id)
        return {
            "summary_sections": [],
            "formatted_text": "",
            "source": "A17-1",
            "completeness": 0.0,
            "message": "未找到 A17-1 底稿，请先创建并完成 A17-1 章节编辑",
        }

    chapter_ids = [ch[0] for ch in _KEY_CHAPTERS]
    rows = await _load_chapter_remarks(db, wp_id, chapter_ids)

    sections: list[dict] = []
    filled = 0
    for ch_id, ch_title in _KEY_CHAPTERS:
        content = (rows.get(ch_id) or "").strip()
        if content:
            filled += 1
        sections.append({"title": ch_title, "content": content, "chapter_id": ch_id})

    kam_text = await _load_kam_summary(db, project_id)
    if kam_text:
        sections.append(
            {
                "title": "关键审计事项（KAM 摘要）",
                "content": kam_text,
                "chapter_id": "A17-2-1",
            }
        )
        filled += 1

    completeness = filled / len(sections) if sections else 0.0
    formatted_parts = [
        f"【{s['title']}】\n{s['content']}" for s in sections if s.get("content")
    ]
    formatted_text = "\n\n".join(formatted_parts)

    return {
        "summary_sections": sections,
        "formatted_text": formatted_text,
        "source": "A17-1",
        "completeness": round(completeness, 2),
        "message": None if formatted_text else "关键章节尚未填写，请先在 A17-1 中编辑",
    }


def build_a18_1_replacements(client_name: str, audit_year: str) -> dict[str, str]:
    """A18-1 模板占位符替换表（信函抬头 + 年度）。"""
    prev_year = str(int(audit_year) - 1) if audit_year.isdigit() else "201X"
    display_name = client_name
    if display_name and "股份有限公司" not in display_name and "公司" not in display_name:
        display_name = f"{display_name}股份有限公司"
    return {
        "××公司": client_name,
        "ABC公司": client_name,
        "XX公司": client_name,
        "XX股份有限公司": display_name,
        "XX年度": f"{audit_year}年度",
        "202X": audit_year,
        "201X": prev_year,
    }


def apply_replacements_to_document(doc: Document, replacements: dict[str, str]) -> None:
    """在 docx 段落/表格中替换占位符。"""
    for para in doc.paragraphs:
        for key, val in replacements.items():
            if key in para.text:
                for run in para.runs:
                    if key in run.text:
                        run.text = run.text.replace(key, val)
    for table in doc.tables:
        for row in table.rows:
            for cell in row.cells:
                for para in cell.paragraphs:
                    for key, val in replacements.items():
                        if key in para.text:
                            for run in para.runs:
                                if key in run.text:
                                    run.text = run.text.replace(key, val)


def _signature_paragraph_index(doc: Document) -> int:
    for i, para in enumerate(doc.paragraphs):
        if "致同" in para.text:
            return i
    return len(doc.paragraphs)


def append_summary_to_document(doc: Document, summary: dict) -> bool:
    """在签名区前插入 A17 生成的小结正文。返回是否插入了内容。"""
    sections = [s for s in summary.get("summary_sections", []) if (s.get("content") or "").strip()]
    if not sections:
        return False

    sig_idx = _signature_paragraph_index(doc)
    anchor = doc.paragraphs[sig_idx]

    blocks: list[str] = ["", "审计情况小结（自 A17-1 自动生成，请复核）"]
    for section in sections:
        blocks.append(f"【{section['title']}】")
        blocks.append(section["content"])

    for text in reversed(blocks):
        anchor.insert_paragraph_before(text)
    return True


async def enrich_a18_1_document(
    doc: Document,
    db: AsyncSession,
    project_id: UUID,
    client_name: str,
    audit_year: str,
) -> dict:
    """A18-1 prefilled-download：替换抬头占位符 + 追加 A17 小结。"""
    apply_replacements_to_document(doc, build_a18_1_replacements(client_name, audit_year))
    summary = await generate_audit_summary(db, project_id)
    summary["summary_injected"] = append_summary_to_document(doc, summary)
    return summary

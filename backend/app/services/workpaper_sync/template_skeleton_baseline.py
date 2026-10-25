"""模板骨架行的「模板原值」基线 —— 判 OO 侧读回的骨架行有没有带用户内容。

spec d2-online-edit-platform-closure · R12。

═══ 为什么需要 ═══

instrumentation 为模板受管区的**每一行**预生成骨架身份 ``GTROW-{template_id}-{row:04d}``
（``ExcelInstrumentationSpec.row_uuid``），而 materialize 永不收敛这些行
（``contracts.is_template_skeleton_identity`` 的 docstring）⇒ 它们以**模板原样**留在每一份
representation 里，OO 侧 extract 必然把它们读回来。

对「HTML store 只装业务行」的 entry（D2：前端自铸 ``dr-*`` 行身份，store 里本无骨架行），
OO→HTML 镜像遇到 base 里没有的骨架行时，「一律新建」与「一律跳过」两条路**都错**：

* 一律新建 ⇒ 模板示例内容（D2-2 的序号 1~12 +「单项计提 / 客户类型组合 组合1…组合X」、
  D2-3 的「其中：」）被当业务行写进 HTML；删了，下次 OO 保存又回来（骨架行永不收敛）。
  D2-3 组合区更直接：骨架行没有 category 信号 ⇒ ``StorePayloadError``，整次镜像失败。
* 一律跳过（``phase5_row_table_sheet`` 的先例 —— 那里 HTML 侧的输入骨架行**本就落库**，
  base 里没有的只剩派生行，跳过是对的）⇒ 用户在 OO 里往模板空行里填的数据**静默**
  不进 HTML（新企业「打开模板直接填」这条最自然的路径整段丢失，两侧数字从此对不上）。

唯一正确的判据是：**这一行的可编辑格与模板原值是否一致**。一致 ⇒ 纯模板内容，跳过；
出现了模板里没有的非空值 ⇒ 用户内容，照常镜像。

═══ 判据的三个刻意选择 ═══

1. **只看「用户写了模板里没有的非空值」**，不看「用户删掉了模板文字」：把示例文字清空
   而没写任何新内容的行，没有可镜像的业务内容 —— 新建它只会在 HTML 里造一条空壳行。
2. **比不了就当用户内容**（保守方向）：骨架身份不在基线里（未知 template_id / 越界行号）、
   字段在契约里找不到列、两侧值无法按 value_type 规范化 —— 一律判「带用户内容」。
   这个方向的最坏后果是多建一行（可见、可删），反方向的最坏后果是静默丢数据。
3. **比较口径复用 ``merge.values_equal``**（与三方合并同一份规范化），不另写一份。
"""

from __future__ import annotations

import io
from dataclasses import dataclass, field
from typing import Any, Callable, Iterable, Mapping, Sequence, Set

from app.services.workpaper_sync.contracts import (
    FieldMode,
    SyncContract,
    ValueType,
    is_template_skeleton_identity,
)
from app.services.workpaper_sync.models import SyncDomainError


class TemplateSkeletonBaselineError(SyncDomainError):
    """模板基线建不起来（模板缺受管 sheet / instrumentation 声明与模板对不上）。"""

    error_code = "template_skeleton_baseline_unavailable"


def _blank(value: Any) -> bool:
    return value is None or (isinstance(value, str) and not value.strip())


def _same_value(value: Any, template_value: Any, value_type: ValueType) -> bool:
    from app.services.workpaper_sync.merge import ValueNormalizationError, values_equal

    try:
        return values_equal(value, template_value, value_type)
    except ValueNormalizationError:
        return False  # 规范化不了 ⇒ 证明不了是模板内容 ⇒ 按用户内容处理（保守方向）


@dataclass(frozen=True)
class SkeletonBaseline:
    """骨架行的模板原值。

    * ``cells``：``{骨架行身份: {列字母: 模板原值}}``（只存非空格）
    * ``columns``：``{(table_key, field_id): (列字母, value_type)}``
    """

    cells: Mapping[str, Mapping[str, Any]] = field(default_factory=dict)
    columns: Mapping[tuple[str, str], tuple[str, ValueType]] = field(default_factory=dict)

    def carries_user_content(
        self, *, table_key: str, row_identity: str, fields: Iterable[tuple[str, Any]]
    ) -> bool:
        """这一行（只给**非受保护**字段）是否带了模板里没有的非空值。"""
        template_row = self.cells.get(row_identity)
        for field_id, value in fields:
            if _blank(value):
                continue
            if template_row is None:
                return True
            column_vt = self.columns.get((table_key, field_id))
            if column_vt is None:
                return True
            column, value_type = column_vt
            if not _same_value(value, template_row.get(column), value_type):
                return True
        return False


def build_skeleton_baseline(
    *,
    template_bytes: bytes,
    specs: Sequence[Any],
    contract: SyncContract,
) -> SkeletonBaseline:
    """从**权威模板字节** + instrumentation 声明 + 契约建基线。

    ``specs`` 是 ``ExcelInstrumentationSpec`` 序列（取 ``managed_sheet`` /
    ``first_data_row`` / ``last_data_row`` / ``managed_last_col`` / ``row_uuid``）——
    骨架身份与模板行号的对应关系**只**从这里来，与 instrumentation 生成身份的规则同源。
    """
    import openpyxl
    from openpyxl.utils import column_index_from_string, get_column_letter

    try:
        wb = openpyxl.load_workbook(io.BytesIO(template_bytes), read_only=True, data_only=False)
    except Exception as exc:  # noqa: BLE001 - 统一翻译成可诊断的领域错误
        raise TemplateSkeletonBaselineError(
            f"权威模板字节无法解析: {type(exc).__name__}: {exc}"
        ) from exc
    cells: dict[str, dict[str, Any]] = {}
    try:
        for spec in specs:
            if spec.managed_sheet not in wb.sheetnames:
                raise TemplateSkeletonBaselineError(
                    f"instrumentation 声明的受管 sheet {spec.managed_sheet!r} 不在模板里"
                    f"（模板 sheet: {wb.sheetnames}）"
                )
            ws = wb[spec.managed_sheet]
            last_col = column_index_from_string(spec.managed_last_col)
            rows = ws.iter_rows(
                min_row=spec.first_data_row,
                max_row=spec.last_data_row,
                min_col=1,
                max_col=last_col,
                values_only=True,
            )
            for offset, row_values in enumerate(rows):
                identity = spec.row_uuid(spec.first_data_row + offset)
                if not is_template_skeleton_identity(identity):
                    raise TemplateSkeletonBaselineError(
                        f"instrumentation 生成的身份 {identity!r} 不是骨架身份形态 —— "
                        "与 contracts.is_template_skeleton_identity 的判据漂移了"
                    )
                cells[identity] = {
                    get_column_letter(index + 1): value
                    for index, value in enumerate(row_values)
                    if value is not None
                }
    finally:
        wb.close()

    columns: dict[tuple[str, str], tuple[str, ValueType]] = {}
    for sheet in contract.sheets:
        for table in sheet.tables:
            for spec_field in table.fields:
                cell = getattr(spec_field, "cell", None)
                if cell is None or not getattr(cell, "column", None):
                    continue
                field_id = str(spec_field.stable_field_key).rsplit("/", 1)[-1]
                columns[(table.table_key, field_id)] = (cell.column, spec_field.value_type)
    return SkeletonBaseline(cells=cells, columns=columns)


def pure_template_skeleton_rows(
    projection: Any,
    *,
    table_key: str,
    base_identities: Set[str],
    baseline: SkeletonBaseline | Callable[[], SkeletonBaseline],
) -> frozenset[str]:
    """本表里「骨架身份 ∧ base 没有 ∧ 不带用户内容」的行身份 —— 镜像应跳过它们。

    只看**非受保护、非公式**字段是否带用户内容（受保护字段本就不回写 store，模板公式的缓存值 /
    OO 重算值也不是用户填的）。只有受保护 / 公式字段的骨架行同样算「不带用户内容」：store 镜像
    视图会把采纳的 OO 重算值放行成不受保护（spec oo-formula-recompute-not-tamper R3.3），这类行
    若不在结果里，镜像会为它新建一条只有公式值的 store 行 —— 正是 R12 要防的骨架行回流。

    ``baseline`` 可传零参可调用对象：没有带可编辑值的候选行（最常见）时**不**读模板。
    """
    prefix = f"{table_key}/"
    per_row: dict[str, list[tuple[str, Any]]] = {}
    for key in projection.stable_keys():
        stable_key = str(key)
        if not stable_key.startswith(prefix):
            continue
        fv = projection.get(key)
        if fv is None:
            continue
        rid = str(getattr(fv, "row_key", None) or "")
        if not rid or rid in base_identities or not is_template_skeleton_identity(rid):
            continue
        fields = per_row.setdefault(rid, [])
        if getattr(fv, "is_protected", False) or getattr(fv, "mode", None) is FieldMode.formula:
            continue
        fields.append((stable_key.rsplit("/", 1)[-1], getattr(fv, "value", None)))
    pure = {rid for rid, fields in per_row.items() if not fields}
    candidates = {rid: fields for rid, fields in per_row.items() if fields}
    if not candidates:
        return frozenset(pure)
    resolved = baseline if isinstance(baseline, SkeletonBaseline) else baseline()
    return frozenset(
        pure
        | {
            rid
            for rid, fields in candidates.items()
            if not resolved.carries_user_content(
                table_key=table_key, row_identity=rid, fields=fields
            )
        }
    )


__all__ = [
    "SkeletonBaseline",
    "TemplateSkeletonBaselineError",
    "build_skeleton_baseline",
    "pure_template_skeleton_rows",
]

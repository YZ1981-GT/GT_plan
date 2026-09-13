"""D4-29: logical customer rows rendered as identity-bearing Excel columns."""
from __future__ import annotations

import copy
import io
import json
from typing import Mapping, Sequence

from openpyxl import load_workbook
from openpyxl.utils import get_column_letter, column_index_from_string
from openpyxl.styles import Protection

from app.services.workpaper_sync.adapters.base import FieldValue, Projection
from app.services.workpaper_sync.definitions import canonical_digest

MANAGED_SHEET = "客户信息检查表D4-29"
SHEET_KEY = "d4-29-managed"
TABLE_KEY = "customer_detail_transposed"
TEMPLATE_ID = "D429"
STORE_ITEM_ID = "D4-29-customers"
IDENTITY_KEY = "id"
HEADER_ROW = 10
FIRST_FIELD_ROW = 11
LAST_FIELD_ROW = 41
FOOTER_ROWS = (42, 46)
STATIC_PROMPT_FIRST_ROW = 48
FIRST_CUSTOMER_COLUMN = "C"
INITIAL_CUSTOMER_COLUMN = "M"
IDENTITY_CARRIER_ROW = 9
IDENTITY_CARRIER_PREFIX = "GT-CUSTOMER-"
# Rows 27 and 32..33 are source-template continuation slots, not new fields.
FIELD_ROWS = {
    "creditCode": 11, "regAddress": 12, "officeAddress": 13, "website": 14,
    "websiteIp": 15, "email": 16, "establishDate": 17, "registeredCapital": 18,
    "bizScope": 19, "headcount": 20, "legalRep": 21, "shareholder1": 22,
    "shareholder2": 23, "shareholder3": 24, "shareholder4": 25, "shareholder5": 26,
    "chairman": 28, "gm": 29, "otherMgmt": 30, "keyHandler": 31,
    "actualController": 34, "isRelated": 35, "isAlsoSupplier": 36,
    "cooperationStart": 37, "hasOverdue": 38, "bizStatus": 39,
    "isBlacklisted": 40, "infoSource": 41,
}
FIELD_KEYS = tuple(FIELD_ROWS)


def _payload(payload):
    if payload is None:
        return []
    if isinstance(payload, (str, bytes, bytearray)):
        payload = json.loads(payload or "[]")
    if isinstance(payload, Mapping):
        payload = payload.get("customers", [])
    if not isinstance(payload, Sequence) or isinstance(payload, (str, bytes, bytearray)):
        raise ValueError("D4-29 customers must be a JSON array")
    out, seen = [], set()
    for raw in payload:
        if not isinstance(raw, Mapping):
            raise ValueError("D4-29 customer must be an object")
        ident = str(raw.get("id", "")).strip()
        if not ident or ident in seen or any(c in ident for c in "/~{}"):
            raise ValueError("D4-29 customer.id must be unique and safe")
        fields = raw.get("fields", {})
        if not isinstance(fields, Mapping):
            raise ValueError("D4-29 customer.fields must be an object")
        out.append({"id": ident, "name": str(raw.get("name", "") or ""), "fields": dict(fields)})
        seen.add(ident)
    return out


def stable_key_for(customer_id, field_key):
    return f"{TABLE_KEY}/{customer_id}/{field_key.lower()}"


def build_store_projection(payload, *, contract, limits=None):
    from app.services.workpaper_sync.excel_extract import StreamingProjectionBudget
    from app.services.workpaper_sync.limits import load_limits
    budget = StreamingProjectionBudget(limits or load_limits())
    values = {}
    customers = _payload(payload)
    for customer in customers:
        budget.add_row(TABLE_KEY)
        for key, value in (("name", customer["name"]), *((k, customer["fields"].get(k, "")) for k in FIELD_KEYS)):
            spec = contract.field_by_stable_key(stable_key_for("{row_uuid}", key))
            stable = stable_key_for(customer["id"], key)
            budget.add_field()
            values[stable] = FieldValue(stable, value, spec.value_type, spec.mode, row_key=customer["id"])
    return Projection(contract_id=contract.contract_id, semantic_version=contract.semantic_version,
                      document_type=contract.document_type, values=values,
                      row_keys={TABLE_KEY: tuple(c["id"] for c in customers)})


def merge_projection_into_store(*, projection, base_payload):
    rows = _payload(base_payload)
    by = {r["id"]: r for r in rows}
    original = set(by)
    order = list(by)
    applied = visited = 0
    keys = {k.lower(): k for k in ("name", *FIELD_KEYS)}
    for stable in projection.stable_keys():
        if not str(stable).startswith(TABLE_KEY + "/"):
            continue
        _, ident, key = str(stable).split("/", 2)
        field = projection.get(stable)
        if field is None or field.is_protected or key not in keys:
            continue
        key = keys[key]
        if ident not in by:
            by[ident] = {"id": ident, "name": "", "fields": {}}
            order.append(ident)
            applied += 1
        target = by[ident] if key == "name" else by[ident]["fields"]
        value = "" if field.value is None else field.value
        visited += 1
        if target.get(key, "") != value:
            target[key] = value
            applied += 1
    ids = list(projection.row_keys[TABLE_KEY]) if TABLE_KEY in projection.row_keys else order
    if len(ids) != len(set(ids)) or any(i not in by for i in ids):
        raise ValueError("D4-29 invalid projection customer identities")
    removed = original - set(ids)
    applied += len(removed) + int(ids != order and not removed)
    return [by[i] for i in ids], applied, visited, removed


def _src(cell):
    return f"源xlsx!{MANAGED_SHEET}!{cell}"


def _parse_field_pointer(key: str, *, row_uuid: str) -> str:
    return f"/customers/{row_uuid}/{'name' if key == 'name' else 'fields/' + key}"


def sheet_payload():
    fields = []
    for key in ("name", *FIELD_KEYS):
        row = HEADER_ROW if key == "name" else FIELD_ROWS[key]
        fields.append({
            "stable_field_key": stable_key_for("{row_uuid}", key),
            "json_pointer": _parse_field_pointer(key, row_uuid="{row_uuid}"),
            "source_ref": _src(f"C{row}:M{row}"),
            "header_source_ref": _src(f"A{row}:B{row}"),
            "cell": {"column": "C", "row_from": "row_identity"},
            "transposed_row": row, "mode": "editable", "value_type": "text",
            "instances": "many",
        })
    return {"sheet_key": SHEET_KEY, "template_id": TEMPLATE_ID, "excel_name": MANAGED_SHEET,
            "tables": [{
                "table_key": TABLE_KEY, "anchor": "C10", "header_rows": 1,
                "layout": "customer_columns", "row_identity": {"kind": "field", "json_pointer": "/customers/*/id"},
                "delete_policy": "tombstone",
                "transposed_columns": {"source_ref": _src("C10:M10"), "identity_row": IDENTITY_CARRIER_ROW,
                                       "identity_prefix": IDENTITY_CARRIER_PREFIX, "identity": "id"},
                "protected_regions": {"label_columns": "A:B", "footer_rows": list(FOOTER_ROWS),
                                      "static_prompt": f"A{STATIC_PROMPT_FIRST_ROW}:XFD1048576"},
                "fields": fields}]}


def mapping_digest_payload():
    return sheet_payload()


def compute_mapping_digest():
    return canonical_digest(mapping_digest_payload())


EXPECTED_MAPPING_DIGEST = "36b1656cf6468789b6e7136f32c6f68f7b6515083e4dd5cbd1609d764560925a"


def assert_mapping_digest():
    got = compute_mapping_digest()
    if got != EXPECTED_MAPPING_DIGEST:
        raise ValueError(f"D4-29 mapping_digest drift: {got} != {EXPECTED_MAPPING_DIGEST}")
    return got


def _copy_column(ws, src, dst):
    source, target = get_column_letter(src), get_column_letter(dst)
    dimension = copy.copy(ws.column_dimensions[source])
    dimension.index = target
    dimension.min = dimension.max = dst
    ws.column_dimensions[target] = dimension
    for row in range(1, ws.max_row + 1):
        a, b = ws.cell(row, src), ws.cell(row, dst)
        b._style = copy.copy(a._style)
        if row < FOOTER_ROWS[0]:
            b.value = a.value
            if a.comment:
                b.comment = copy.copy(a.comment)
            if a.hyperlink:
                b.hyperlink = copy.copy(a.hyperlink)
    for merged in list(ws.merged_cells.ranges):
        if merged.min_col == merged.max_col == src and merged.max_row < FOOTER_ROWS[0]:
            ws.merge_cells(start_row=merged.min_row, end_row=merged.max_row, start_column=dst, end_column=dst)


def materialize_transposed_workbook(workbook_bytes: bytes, payload, *, sheet_name=MANAGED_SHEET) -> bytes:
    wb = load_workbook(io.BytesIO(workbook_bytes))
    ws = wb[sheet_name]
    customers = _payload(payload)
    start = column_index_from_string(FIRST_CUSTOMER_COLUMN)
    last = column_index_from_string(INITIAL_CUSTOMER_COLUMN)
    if start + len(customers) - 1 > 16384:
        raise ValueError("D4-29 exceeds Excel column limit")
    for i in range(len(customers)):
        col = start + i
        if col > last:
            _copy_column(ws, last, col)
    for col in range(start, ws.max_column + 1):
        for row in (IDENTITY_CARRIER_ROW, HEADER_ROW, *FIELD_ROWS.values()):
            ws.cell(row, col).value = None
    ws.row_dimensions[IDENTITY_CARRIER_ROW].hidden = True
    for i, customer in enumerate(customers):
        col = start + i
        ws.cell(IDENTITY_CARRIER_ROW, col).value = IDENTITY_CARRIER_PREFIX + customer["id"]
        ws.cell(IDENTITY_CARRIER_ROW, col).protection = Protection(locked=True)
        for key in ("name", *FIELD_KEYS):
            row = HEADER_ROW if key == "name" else FIELD_ROWS[key]
            cell = ws.cell(row, col)
            cell.value = customer["name"] if key == "name" else customer["fields"].get(key, "")
            cell.data_type = "s" if isinstance(cell.value, str) else cell.data_type
            cell.protection = Protection(locked=False)
    ws.protection.sheet = True
    out = io.BytesIO()
    wb.save(out)
    # Keep unrelated sheets, caches, relationships and unsupported OOXML untouched.
    import zipfile
    from xml.etree import ElementTree as ET
    with zipfile.ZipFile(io.BytesIO(workbook_bytes)) as original, zipfile.ZipFile(io.BytesIO(out.getvalue())) as edited:
        ns = {"m": "http://schemas.openxmlformats.org/spreadsheetml/2006/main"}
        rel_ns = "{http://schemas.openxmlformats.org/officeDocument/2006/relationships}id"
        sheets = ET.fromstring(original.read("xl/workbook.xml"))
        rid = next(s.attrib[rel_ns] for s in sheets.findall("m:sheets/m:sheet", ns) if s.attrib["name"] == sheet_name)
        rels = ET.fromstring(original.read("xl/_rels/workbook.xml.rels"))
        target = next(r.attrib["Target"] for r in rels if r.attrib["Id"] == rid)
        part = target.lstrip("/") if target.startswith("/") else "xl/" + target
        edited_sheets = ET.fromstring(edited.read("xl/workbook.xml"))
        edited_rid = next(s.attrib[rel_ns] for s in edited_sheets.findall("m:sheets/m:sheet", ns) if s.attrib["name"] == sheet_name)
        edited_rels = ET.fromstring(edited.read("xl/_rels/workbook.xml.rels"))
        edited_target = next(r.attrib["Target"] for r in edited_rels if r.attrib["Id"] == edited_rid)
        edited_part = edited_target.lstrip("/") if edited_target.startswith("/") else "xl/" + edited_target
        result = io.BytesIO()
        with zipfile.ZipFile(result, "w", zipfile.ZIP_DEFLATED) as archive:
            for info in original.infolist():
                data = edited.read(edited_part) if info.filename == part else (
                    edited.read("xl/styles.xml") if info.filename == "xl/styles.xml" else original.read(info.filename))
                archive.writestr(info, data)
        return result.getvalue()


def extract_transposed_workbook(workbook_bytes: bytes, *, sheet_name=MANAGED_SHEET):
    ws = load_workbook(io.BytesIO(workbook_bytes), data_only=False)[sheet_name]
    result = []
    for col in range(column_index_from_string(FIRST_CUSTOMER_COLUMN), ws.max_column + 1):
        raw = ws.cell(IDENTITY_CARRIER_ROW, col).value
        if not raw:
            continue
        if not isinstance(raw, str) or not raw.startswith(IDENTITY_CARRIER_PREFIX):
            raise ValueError("D4-29 invalid customer identity carrier")
        def value(row):
            cell = ws.cell(row, col)
            if cell.data_type == "f":
                raise ValueError("D4-29 text field cannot contain a formula")
            return "" if cell.value is None else cell.value
        result.append({"id": raw[len(IDENTITY_CARRIER_PREFIX):], "name": value(HEADER_ROW),
                       "fields": {key: value(row) for key, row in FIELD_ROWS.items()}})
    return _payload(result)


def row_oriented_sheets(contract):
    """Exclude explicitly dispatched column layouts from row-table instrumentation."""
    excluded = {s["sheet_key"] for s in contract.canonical_payload.get("sheets", [])
                if any(t.get("layout") == "customer_columns" for t in s.get("tables", []))}
    return [s for s in contract.sheets if s.sheet_key not in excluded]


def is_enabled(contract):
    return contract.contract_id == "d4.revenue_detail" and any(s.sheet_key == SHEET_KEY for s in contract.sheets)


def materialize_file(output, projection, contract):
    if not is_enabled(contract) or TABLE_KEY not in projection.row_keys:
        return
    customers, _, _, _ = merge_projection_into_store(projection=projection, base_payload=[])
    output.write_bytes(materialize_transposed_workbook(output.read_bytes(), customers))


def extract_file(artifact, contract):
    return build_store_projection(extract_transposed_workbook(artifact.read_bytes()), contract=contract)

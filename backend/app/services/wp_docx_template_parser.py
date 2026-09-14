"""Word 模板结构解析器 — 提取 docx 占位符、段落、表格结构.

为 word-template 类型底稿（25 个 wp_code）提供统一的模板解析能力：
- 识别 ${field_id} / ${field_id:label} 新格式占位符
- 识别 ××公司 / XX公司 / 202X年 等 legacy 中文标记
- 提取段落结构（标题层级、样式）和表格结构
- 推断字段数据类型（text/date/textarea/number）
- mtime 缓存（与 wp_code_overrides 相同模式）
"""

from __future__ import annotations

import logging
import os
import re
import unicodedata
from dataclasses import dataclass, field
from datetime import datetime, timezone
from pathlib import Path

from docx import Document

logger = logging.getLogger(__name__)


# ─── 数据类 ───────────────────────────────────────────────────────────────────


@dataclass
class PlaceholderDef:
    """模板中的占位符定义."""

    field_id: str  # 唯一标识，如 "entity_name"
    label: str  # 显示标签，如 "被审计单位"
    data_type: str  # "text" | "date" | "textarea" | "number"
    default_value: str  # 模板中原始文本
    position: dict  # {"paragraph_index": int} 或 {"table_index": int, "row": int, "col": int}
    pattern: str  # 原始占位符模式，如 "${entity_name}" 或 "××公司"


@dataclass
class ParagraphDef:
    """段落结构定义."""

    index: int
    text: str
    style: str  # "Heading 1" | "Normal" | etc.
    heading_level: int  # 0=非标题, 1~6=对应级别
    placeholder_ids: list[str] = field(default_factory=list)


@dataclass
class TableDef:
    """表格结构定义."""

    index: int
    rows: list[list[str]] = field(default_factory=list)  # 每个单元格的文本
    placeholder_ids: list[str] = field(default_factory=list)


@dataclass
class TemplateStructure:
    """模板解析完整输出."""

    paragraphs: list[ParagraphDef] = field(default_factory=list)
    tables: list[TableDef] = field(default_factory=list)
    placeholders: list[PlaceholderDef] = field(default_factory=list)
    metadata: dict = field(default_factory=dict)  # {"template_name", "wp_code", "last_parsed_at"}


# ─── 占位符正则 ────────────────────────────────────────────────────────────────

# 新格式: ${field_id} 或 ${field_id:显示标签}
_NEW_PLACEHOLDER_RE = re.compile(r"\$\{([a-zA-Z_][a-zA-Z0-9_]*)(?::([^}]+))?\}")

# Legacy 中文标记
_LEGACY_PATTERNS: list[tuple[str, str, str]] = [
    # (regex_pattern, field_id, label)
    (r"××公司", "entity_name", "被审计单位"),
    (r"XX公司", "entity_name", "被审计单位"),
    (r"×{2,}公司", "entity_name", "被审计单位"),
    (r"X{2,}公司", "entity_name", "被审计单位"),
    (r"202X年\d{1,2}月\d{1,2}日", "report_date", "报告日期"),
    (r"202X年XX月XX日", "report_date", "报告日期"),
    (r"20XX年XX月XX日", "report_date", "报告日期"),
    (r"202X年\s*\d{0,2}\s*月\s*\d{0,2}\s*日", "report_date", "报告日期"),
    (r"202X年", "audit_year", "审计年度"),
    (r"20XX年", "audit_year", "审计年度"),
    (r"××", "placeholder_generic", "待填内容"),
]

# 编译 legacy 正则
_LEGACY_COMPILED: list[tuple[re.Pattern, str, str]] = [
    (re.compile(pat), fid, label) for pat, fid, label in _LEGACY_PATTERNS
]

# 日期相关关键词
_DATE_KEYWORDS = re.compile(r"日期|年月日|签署日|报告日")

# 数字相关关键词
_NUMBER_KEYWORDS = re.compile(r"金额|万元|元|数量|比例|百分比|%")


# ─── 辅助函数 ──────────────────────────────────────────────────────────────────


def _slugify(text: str) -> str:
    """将中文/英文标签转为 snake_case field_id."""
    # 移除标点
    text = re.sub(r"[^\w\s]", "", text, flags=re.UNICODE)
    # 中文直接用拼音简写或保留（这里简化为下划线连接）
    text = text.strip()
    if not text:
        return "field"
    # 如果全英文，直接 snake_case
    if text.isascii():
        return re.sub(r"\s+", "_", text.lower())
    # 中文：每个词用下划线连接
    return re.sub(r"\s+", "_", text)


def _dedupe_field_id(base_id: str, existing_ids: set[str]) -> str:
    """生成唯一 field_id，重复时追加计数器."""
    if base_id not in existing_ids:
        return base_id
    counter = 2
    while f"{base_id}_{counter}" in existing_ids:
        counter += 1
    return f"{base_id}_{counter}"


def _infer_data_type(text: str, context_paragraph: str = "") -> str:
    """根据上下文推断字段数据类型."""
    combined = text + context_paragraph
    if _DATE_KEYWORDS.search(combined):
        return "date"
    if _NUMBER_KEYWORDS.search(combined):
        return "number"
    # 长段落上下文 → textarea
    if len(context_paragraph) > 100:
        return "textarea"
    return "text"


def _get_heading_level(paragraph) -> int:
    """从段落获取标题级别. 0=非标题, 1~6=对应级别."""
    style_name = paragraph.style.name if paragraph.style else ""
    # python-docx style name 格式: "Heading 1", "Heading 2", etc.
    match = re.match(r"Heading\s+(\d+)", style_name, re.IGNORECASE)
    if match:
        level = int(match.group(1))
        return min(level, 6)
    # 中文标题样式
    if "标题" in style_name:
        digit_match = re.search(r"(\d+)", style_name)
        if digit_match:
            return min(int(digit_match.group(1)), 6)
        return 1
    return 0


def _extract_placeholders_from_text(
    text: str,
    position: dict,
    existing_ids: set[str],
    context_paragraph: str = "",
) -> list[PlaceholderDef]:
    """从文本中提取所有占位符."""
    placeholders: list[PlaceholderDef] = []

    # 1. 新格式 ${field_id} / ${field_id:label}
    for match in _NEW_PLACEHOLDER_RE.finditer(text):
        raw_id = match.group(1)
        label = match.group(2) or raw_id
        field_id = _dedupe_field_id(raw_id, existing_ids)
        existing_ids.add(field_id)
        data_type = _infer_data_type(label, context_paragraph)
        placeholders.append(
            PlaceholderDef(
                field_id=field_id,
                label=label,
                data_type=data_type,
                default_value=match.group(0),
                position=position.copy(),
                pattern=match.group(0),
            )
        )

    # 2. Legacy 中文标记（仅在新格式未覆盖的文本区域中查找）
    # 移除已匹配的新格式占位符后再查找 legacy
    remaining_text = _NEW_PLACEHOLDER_RE.sub("", text)
    for regex, base_field_id, label in _LEGACY_COMPILED:
        for match in regex.finditer(remaining_text):
            field_id = _dedupe_field_id(base_field_id, existing_ids)
            existing_ids.add(field_id)
            data_type = _infer_data_type(label, context_paragraph)
            placeholders.append(
                PlaceholderDef(
                    field_id=field_id,
                    label=label,
                    data_type=data_type,
                    default_value=match.group(0),
                    position=position.copy(),
                    pattern=match.group(0),
                )
            )

    return placeholders


# ─── 公开 API ─────────────────────────────────────────────────────────────────


def parse_template(file_path: str) -> TemplateStructure:
    """解析 docx 模板文件，返回完整结构化数据.

    Args:
        file_path: docx 模板文件路径

    Returns:
        TemplateStructure 包含段落、表格、占位符列表

    Raises:
        FileNotFoundError: 文件不存在
        ValueError: 文件不是有效的 docx 格式
    """
    path = Path(file_path)

    if not path.exists():
        raise FileNotFoundError(f"模板文件不存在: {file_path}")

    try:
        doc = Document(str(path))
    except Exception as exc:
        raise ValueError(f"Invalid docx file: {file_path}") from exc

    paragraphs: list[ParagraphDef] = []
    tables: list[TableDef] = []
    all_placeholders: list[PlaceholderDef] = []
    existing_ids: set[str] = set()

    # ─── 解析段落 ─────────────────────────────────────────────────────────
    for idx, para in enumerate(doc.paragraphs):
        text = para.text.strip()
        style_name = para.style.name if para.style else "Normal"
        heading_level = _get_heading_level(para)

        # 提取段落内占位符
        position = {"paragraph_index": idx}
        para_placeholders = _extract_placeholders_from_text(
            text, position, existing_ids, context_paragraph=text
        )
        placeholder_ids = [p.field_id for p in para_placeholders]
        all_placeholders.extend(para_placeholders)

        paragraphs.append(
            ParagraphDef(
                index=idx,
                text=text,
                style=style_name,
                heading_level=heading_level,
                placeholder_ids=placeholder_ids,
            )
        )

    # ─── 解析表格 ─────────────────────────────────────────────────────────
    for tbl_idx, table in enumerate(doc.tables):
        rows_data: list[list[str]] = []
        table_placeholder_ids: list[str] = []

        for row_idx, row in enumerate(table.rows):
            row_cells: list[str] = []
            for col_idx, cell in enumerate(row.cells):
                cell_text = cell.text.strip()
                row_cells.append(cell_text)

                # 提取单元格内占位符
                position = {"table_index": tbl_idx, "row": row_idx, "col": col_idx}
                cell_placeholders = _extract_placeholders_from_text(
                    cell_text, position, existing_ids, context_paragraph=cell_text
                )
                for p in cell_placeholders:
                    table_placeholder_ids.append(p.field_id)
                all_placeholders.extend(cell_placeholders)

            rows_data.append(row_cells)

        tables.append(
            TableDef(
                index=tbl_idx,
                rows=rows_data,
                placeholder_ids=table_placeholder_ids,
            )
        )

    # ─── 构建元数据 ───────────────────────────────────────────────────────
    template_name = path.stem
    metadata = {
        "template_name": template_name,
        "wp_code": "",  # 由调用方填充
        "last_parsed_at": datetime.now(timezone.utc).isoformat(),
    }

    structure = TemplateStructure(
        paragraphs=paragraphs,
        tables=tables,
        placeholders=all_placeholders,
        metadata=metadata,
    )

    logger.info(
        "解析模板完成: %s, 段落=%d, 表格=%d, 占位符=%d",
        path.name,
        len(paragraphs),
        len(tables),
        len(all_placeholders),
    )

    return structure


# ─── 缓存管理 ─────────────────────────────────────────────────────────────────

_TEMPLATE_CACHE: dict[str, tuple[float, TemplateStructure]] = {}


def get_cached_structure(file_path: str, wp_code: str) -> TemplateStructure:
    """获取带 mtime 缓存的模板结构.

    同一文件未修改时返回缓存对象（identity）；mtime 变化时重新解析。

    Args:
        file_path: docx 模板文件路径
        wp_code: 底稿编码（填入 metadata）

    Returns:
        TemplateStructure

    Raises:
        FileNotFoundError: 文件不存在
        ValueError: 文件不是有效的 docx 格式
    """
    path = Path(file_path)
    if not path.exists():
        raise FileNotFoundError(f"模板文件不存在: {file_path}")

    current_mtime = os.path.getmtime(path)
    cache_key = str(path.resolve())

    if cache_key in _TEMPLATE_CACHE:
        cached_mtime, cached_structure = _TEMPLATE_CACHE[cache_key]
        if cached_mtime == current_mtime:
            return cached_structure

    # 解析并缓存
    structure = parse_template(file_path)
    structure.metadata["wp_code"] = wp_code
    _TEMPLATE_CACHE[cache_key] = (current_mtime, structure)
    return structure


def format_placeholder_summary(structure: TemplateStructure) -> str:
    """格式化占位符摘要（用于 round-trip 验证）.

    输出格式：每行一个占位符，格式为 field_id|label|data_type|pattern
    """
    lines: list[str] = []
    for p in structure.placeholders:
        lines.append(f"{p.field_id}|{p.label}|{p.data_type}|{p.pattern}")
    return "\n".join(lines)


def invalidate_cache(file_path: str | None = None) -> None:
    """清除模板缓存.

    Args:
        file_path: 指定路径则只清该条目；None 清除全部缓存
    """
    global _TEMPLATE_CACHE
    if file_path is None:
        _TEMPLATE_CACHE.clear()
        logger.info("已清除全部模板解析缓存")
    else:
        cache_key = str(Path(file_path).resolve())
        if cache_key in _TEMPLATE_CACHE:
            del _TEMPLATE_CACHE[cache_key]
            logger.info("已清除模板缓存: %s", file_path)

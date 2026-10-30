"""批量建项 Excel 协议、模板和受限解析。

数据 sheet 是唯一导入源；示例、选项和元数据分别放独立 sheet，避免示例被误建项。
"""
from __future__ import annotations

import hashlib
import json
import re
import zipfile
from dataclasses import dataclass
from io import BytesIO
from uuid import UUID

from fastapi import HTTPException
from openpyxl import Workbook, load_workbook
from openpyxl.comments import Comment
from openpyxl.styles import Font, PatternFill, Protection
from openpyxl.utils import get_column_letter
from openpyxl.worksheet.datavalidation import DataValidation
from openpyxl.workbook.defined_name import DefinedName
from pydantic import BaseModel

from app.services.group_relation import effective_parent_code
from app.services.uscc_validator import validate_uscc

TEMPLATE_VERSION = "2"
MAX_IMPORT_ROWS = 500
MAX_UPLOAD_BYTES = 5 * 1024 * 1024
MAX_UNCOMPRESSED_BYTES = 24 * 1024 * 1024
MAX_ZIP_MEMBERS = 1000

BASE_COLUMNS = [
    "客户名称",
    "企业代码(USCC)",
    "项目简称",
    "审计年度",
    "项目类型",
    "会计准则",
    "报表类型",
    "上级企业代码(parent)",
    "最终控制方代码(ultimate)",
    "与上级关系(relation)",
]
EXTENDED_COLUMNS = BASE_COLUMNS + [
    "操作类型(action)",
    "项目ID(project_id)",
    "关系基线摘要(base_hash)",
]

PROJECT_TYPE_MAP = {
    "年报审计": "annual",
    "专项审计": "special",
    "IPO审计": "ipo",
    "内控审计": "internal_control",
    "验资": "capital_verification",
    "税审": "tax_audit",
}
ACCOUNTING_STANDARD_MAP = {
    "企业会计准则": "enterprise",
    "小企业会计准则": "small_enterprise",
    "金融企业会计准则": "financial",
    "政府会计准则": "government",
}
REPORT_SCOPE_MAP = {"单户": "standalone", "合并": "consolidated"}
RELATION_MAP = {"子公司": "subsidiary", "分公司": "branch"}
ACTION_MAP = {"新建": "create", "更新集团关系": "update", "跳过": "skip"}

PROJECT_TYPE_REVERSE = {value: key for key, value in PROJECT_TYPE_MAP.items()}
ACCOUNTING_STANDARD_REVERSE = {value: key for key, value in ACCOUNTING_STANDARD_MAP.items()}
REPORT_SCOPE_REVERSE = {value: key for key, value in REPORT_SCOPE_MAP.items()}
RELATION_REVERSE = {value: key for key, value in RELATION_MAP.items()}
ACTION_REVERSE = {value: key for key, value in ACTION_MAP.items()}
_FORMULA_PREFIXES = ("=", "+", "-", "@")


def _append_text_safe(sheet, values: list | tuple) -> None:
    """把用户可控文本写成字符串单元格，阻止 Excel 将公式前缀解释为公式。"""
    sheet.append(values)
    for cell in sheet[sheet.max_row]:
        if isinstance(cell.value, str) and cell.value.startswith(_FORMULA_PREFIXES):
            cell.data_type = "s"


EXAMPLE_ROWS = [
    ["示例集团有限公司", "91110000100000000R", "示例集团", 2025, "年报审计", "企业会计准则", "合并", "", "91110000100000000R", "", "新建", "", ""],
    ["示例区域控股有限公司", "911100002000000005", "区域控股", 2025, "年报审计", "企业会计准则", "单户", "91110000100000000R", "91110000100000000R", "子公司", "新建", "", ""],
    ["示例子公司有限公司", "91110000300000000G", "示例子公司", 2025, "年报审计", "企业会计准则", "单户", "911100002000000005", "91110000100000000R", "子公司", "新建", "", ""],
    ["示例集团有限公司上海分公司", "91110000400000000U", "上海分公司", 2025, "年报审计", "企业会计准则", "单户", "91110000100000000R", "91110000100000000R", "分公司", "新建", "", ""],
]

class ParsedRow(BaseModel):
    row_number: int = 0
    client_name: str = ""
    company_code: str = ""
    short_name: str = ""
    audit_year: int | None = None
    project_type: str | None = None
    accounting_standard: str | None = None
    report_scope: str = "standalone"
    parent_company_code: str = ""
    parent_company_name: str = ""
    ultimate_company_code: str = ""
    ultimate_company_name: str = ""
    relation_to_parent: str | None = None
    action: str = "create"
    project_id: UUID | None = None
    base_hash: str = ""


@dataclass(frozen=True)
class LoadedWorkbook:
    rows: list[tuple]
    template_version: str | None
    mode: str | None
    root_project_id: UUID | None
    legacy: bool
    file_sha256: str


def is_blank_row(row: tuple | list) -> bool:
    return not row or all(value is None or str(value).strip() == "" for value in row)


def _reject_unsafe_archive(file_bytes: bytes) -> None:
    if not file_bytes:
        raise HTTPException(status_code=400, detail="上传文件为空")
    if len(file_bytes) > MAX_UPLOAD_BYTES:
        raise HTTPException(status_code=413, detail="文件超过 5MB，请精简后重新上传")
    try:
        with zipfile.ZipFile(BytesIO(file_bytes)) as archive:
            infos = archive.infolist()
            if len(infos) > MAX_ZIP_MEMBERS:
                raise HTTPException(status_code=413, detail="工作簿内部文件数量异常")
            total = sum(max(info.file_size, 0) for info in infos)
            if total > MAX_UNCOMPRESSED_BYTES:
                raise HTTPException(status_code=413, detail="工作簿解压后内容过大")
            for info in infos:
                if info.compress_size and info.file_size > 1024 * 1024:
                    if info.file_size / info.compress_size > 500:
                        raise HTTPException(status_code=413, detail="工作簿压缩比例异常")
    except HTTPException:
        raise
    except (zipfile.BadZipFile, OSError):
        raise HTTPException(status_code=400, detail="文件格式不正确，请上传 .xlsx 文件")


def _read_meta(workbook) -> tuple[str | None, str | None, UUID | None]:
    """读取受模板协议保护的版本、场景与可选集团根标识。"""
    if "_GT_META" not in workbook.sheetnames:
        return None, None, None
    values: dict[str, str] = {}
    for key, value in workbook["_GT_META"].iter_rows(
        min_row=1, max_row=3, min_col=1, max_col=2, values_only=True,
    ):
        if key is None:
            continue
        values[str(key).strip()] = "" if value is None else str(value).strip()
    version = values.get("template_version") or None
    mode = values.get("mode") or None
    root_text = values.get("root_project_id") or ""
    if not version or mode not in {"project", "consolidation"}:
        raise HTTPException(status_code=400, detail="模板元数据不完整或场景无效，请重新下载模板")
    if root_text and mode != "consolidation":
        raise HTTPException(status_code=400, detail="普通建项模板不能携带集团根标识，请重新下载模板")
    try:
        root_project_id = UUID(root_text) if root_text else None
    except (TypeError, ValueError) as exc:
        raise HTTPException(status_code=400, detail="模板集团根标识无效，请重新从当前集团树导出") from exc
    return version, mode, root_project_id


def load_data_sheet(file_bytes: bytes) -> LoadedWorkbook:
    """严格读取数据 sheet；兼容旧 9/10 列，但不猜 active sheet。"""
    _reject_unsafe_archive(file_bytes)
    try:
        workbook = load_workbook(BytesIO(file_bytes), read_only=True, data_only=True)
    except Exception as exc:
        raise HTTPException(status_code=400, detail="文件格式不正确，请使用标准建项模板") from exc
    if "数据" not in workbook.sheetnames:
        raise HTTPException(status_code=400, detail="缺少“数据”sheet，请使用标准建项模板")
    sheet = workbook["数据"]
    raw_headers = next(sheet.iter_rows(min_row=1, max_row=1, values_only=True), ())
    headers = [str(value).strip() if value is not None else "" for value in raw_headers]
    nonempty_count = len(headers)
    while nonempty_count and not headers[nonempty_count - 1]:
        nonempty_count -= 1
    headers = headers[:nonempty_count]
    if headers not in (BASE_COLUMNS[:9], BASE_COLUMNS, EXTENDED_COLUMNS):
        raise HTTPException(status_code=400, detail="模板表头已被修改或列顺序不正确，请重新下载模板")
    if (sheet.max_row or 0) > MAX_IMPORT_ROWS + 100:
        raise HTTPException(status_code=413, detail=f"数据 sheet 超过允许范围（最多 {MAX_IMPORT_ROWS} 行）")
    rows: list[tuple] = []
    max_col = len(EXTENDED_COLUMNS)
    for row in sheet.iter_rows(min_row=2, max_col=max_col, values_only=True):
        if is_blank_row(row):
            continue
        rows.append(tuple(row))
        if len(rows) > MAX_IMPORT_ROWS:
            raise HTTPException(status_code=413, detail=f"文件超过最大行数限制（{MAX_IMPORT_ROWS} 行）")
    version, template_mode, root_project_id = _read_meta(workbook)
    if version is not None and version != TEMPLATE_VERSION:
        raise HTTPException(
            status_code=400,
            detail=f"模板版本 {version} 不受支持，请重新下载当前版本模板",
        )
    return LoadedWorkbook(
        rows=rows,
        template_version=version,
        mode=template_mode,
        root_project_id=root_project_id,
        legacy=headers != EXTENDED_COLUMNS,
        file_sha256=hashlib.sha256(file_bytes).hexdigest(),
    )

def _cell(row: tuple | list, index: int) -> str:
    if index >= len(row) or row[index] is None:
        return ""
    return str(row[index]).strip()


def _parse_year(raw: object, errors: list[str]) -> int | None:
    if raw is None or str(raw).strip() == "":
        errors.append("审计年度为必填项")
        return None
    if isinstance(raw, bool):
        errors.append("审计年度必须为4位整数")
        return None
    if isinstance(raw, int):
        year = raw
    elif isinstance(raw, float):
        if not raw.is_integer():
            errors.append("审计年度必须为4位整数，不能包含小数")
            return None
        year = int(raw)
    else:
        text = str(raw).strip()
        if not re.fullmatch(r"\d{4}", text):
            errors.append("审计年度必须为4位整数")
            return None
        year = int(text)
    if year < 2000 or year > 2100:
        errors.append("审计年度须在 2000-2100 之间")
    return year


def parse_row(row: tuple | list, *, row_number: int = 0) -> tuple[ParsedRow, list[str]]:
    """解析一行；skip 行只校验动作，避免用户必须修复明确跳过的数据。"""
    errors: list[str] = []
    action_text = _cell(row, 10)
    project_id_text = _cell(row, 11)
    base_hash = _cell(row, 12).lower()
    action = ACTION_MAP.get(action_text) if action_text else ("update" if project_id_text else "create")
    if action is None:
        errors.append(f"操作类型无效：{action_text}（只能填 新建、更新集团关系 或 跳过）")
        action = "create"
    project_id: UUID | None = None
    if project_id_text:
        try:
            project_id = UUID(project_id_text)
        except ValueError:
            errors.append("项目ID格式不正确")
    if action == "update" and project_id is None:
        errors.append("更新集团关系必须提供项目ID")
    if base_hash and not re.fullmatch(r"[0-9a-f]{64}", base_hash):
        errors.append("关系基线摘要格式不正确")
    if action == "skip":
        return ParsedRow(row_number=row_number, action=action, project_id=project_id, base_hash=base_hash), errors

    client_name = _cell(row, 0)
    company_code = _cell(row, 1)
    short_name = _cell(row, 2)
    project_type_cn = _cell(row, 4)
    standard_cn = _cell(row, 5)
    scope_cn = _cell(row, 6)
    parent_code = _cell(row, 7)
    ultimate_code = _cell(row, 8)
    relation_cn = _cell(row, 9)

    if not client_name:
        errors.append("客户名称为必填项")
    if not company_code:
        errors.append("企业代码为必填项")
    elif not validate_uscc(company_code)[0]:
        errors.append(validate_uscc(company_code)[1] or "企业代码格式错误")
    if not short_name:
        errors.append("项目简称为必填项")
    for label, code in (("上级企业代码", parent_code), ("最终控制方代码", ultimate_code)):
        if code:
            valid, reason = validate_uscc(code)
            if not valid:
                errors.append(f"{label}格式错误：{reason or 'USCC 格式不合规'}")

    relation = RELATION_MAP.get(relation_cn) if relation_cn else None
    if relation_cn and relation is None:
        errors.append(f"与上级关系无效：{relation_cn}（只能填 子公司 或 分公司）")
    if effective_parent_code(company_code, parent_code) is None:
        relation = None
    year = _parse_year(row[3] if len(row) > 3 else None, errors)

    project_type = PROJECT_TYPE_MAP.get(project_type_cn)
    if not project_type_cn:
        errors.append("项目类型为必填项")
    elif project_type is None:
        errors.append(f"项目类型无效：{project_type_cn}")
    standard = ACCOUNTING_STANDARD_MAP.get(standard_cn)
    if not standard_cn:
        errors.append("会计准则为必填项")
    elif standard is None:
        errors.append(f"会计准则无效：{standard_cn}")
    scope = REPORT_SCOPE_MAP.get(scope_cn, "standalone") if scope_cn else "standalone"
    if scope_cn and scope_cn not in REPORT_SCOPE_MAP:
        errors.append(f"报表类型无效：{scope_cn}")

    return ParsedRow(
        row_number=row_number, client_name=client_name, company_code=company_code,
        short_name=short_name, audit_year=year, project_type=project_type,
        accounting_standard=standard, report_scope=scope,
        parent_company_code=parent_code, ultimate_company_code=ultimate_code,
        relation_to_parent=relation, action=action, project_id=project_id,
        base_hash=base_hash,
    ), errors

_REQUIRED_FILL = PatternFill("solid", fgColor="FFF2CC")
_HEADER_FILL = PatternFill("solid", fgColor="E4D3F5")
_HEADER_COMMENTS = {
    1: "必填：被审计单位全称",
    2: "必填：18位统一社会信用代码",
    3: "必填：项目简称",
    4: "必填：4位整数年度",
    5: "必填：请从下拉选项选择",
    6: "必填：请从下拉选项选择",
    7: "选填：默认单户",
    8: "选填：直接上级企业的统一社会信用代码",
    9: "选填：最终控制方统一社会信用代码",
    10: "有上级时选填，留空按企业名称判断",
    11: "空白模板默认新建；现有树导出后可选更新集团关系或跳过",
    12: "系统字段：新建行请留空",
    13: "系统字段：请勿修改",
}


def _add_list_validation(sheet, column: int, range_name: str) -> None:
    letter = get_column_letter(column)
    validation = DataValidation(
        type="list",
        formula1=f"={range_name}",
        allow_blank=column not in (5, 6),
        showErrorMessage=True,
        errorTitle="选项不正确",
        error="请从下拉选项中选择",
    )
    validation.add(f"{letter}2:{letter}{MAX_IMPORT_ROWS + 1}")
    sheet.add_data_validation(validation)


def _add_options_and_validations(workbook, data) -> None:
    options = workbook.create_sheet("_GT_OPTIONS")
    specs = [
        (5, list(PROJECT_TYPE_MAP)),
        (6, list(ACCOUNTING_STANDARD_MAP)),
        (7, list(REPORT_SCOPE_MAP)),
        (10, list(RELATION_MAP)),
        (11, list(ACTION_MAP)),
    ]
    for option_column, (target_column, values) in enumerate(specs, 1):
        for row, value in enumerate(values, 1):
            options.cell(row=row, column=option_column, value=value)
        option_letter = get_column_letter(option_column)
        range_name = f"GT_BATCH_OPTIONS_{target_column}"
        workbook.defined_names.add(DefinedName(
            range_name,
            attr_text=f"'_GT_OPTIONS'!${option_letter}$1:${option_letter}${len(values)}",
        ))
        _add_list_validation(data, target_column, range_name)
    options.sheet_state = "hidden"


def _style_data_sheet(sheet) -> None:
    sheet.freeze_panes = "A2"
    sheet.auto_filter.ref = f"A1:M{MAX_IMPORT_ROWS + 1}"
    widths = [28, 23, 18, 12, 16, 20, 12, 27, 29, 23, 22, 38, 68]
    for index, width in enumerate(widths, 1):
        sheet.column_dimensions[get_column_letter(index)].width = width
    for index, cell in enumerate(sheet[1], 1):
        cell.font = Font(bold=True)
        cell.fill = _HEADER_FILL
        cell.protection = Protection(locked=True)
        cell.comment = Comment(_HEADER_COMMENTS[index], "致同审计平台")
    for row in sheet.iter_rows(min_row=2, max_row=MAX_IMPORT_ROWS + 1, min_col=1, max_col=13):
        for cell in row:
            cell.protection = Protection(locked=False)
        row[11].protection = Protection(locked=True)
        row[12].protection = Protection(locked=True)
        for index in (1, 2, 3, 4, 5, 6):
            row[index - 1].fill = _REQUIRED_FILL
    sheet.column_dimensions["L"].hidden = True
    sheet.column_dimensions["M"].hidden = True
    sheet.protection.sheet = True
    sheet.protection.selectLockedCells = False
    sheet.protection.selectUnlockedCells = True


def generate_template(*, mode: str = "project") -> BytesIO:
    if mode not in ("project", "consolidation"):
        raise HTTPException(status_code=422, detail="批量建项场景无效")
    workbook = Workbook()
    data = workbook.active
    data.title = "数据"
    data.append(EXTENDED_COLUMNS)
    _style_data_sheet(data)

    examples = workbook.create_sheet("填写示例")
    examples.append(EXTENDED_COLUMNS)
    for row in EXAMPLE_ROWS:
        examples.append(row)
    examples.freeze_panes = "A2"
    examples.sheet_properties.tabColor = "F4B183"

    notes = workbook.create_sheet("说明事项")
    notes.append(["字段", "填写规则"])
    notes_rows = [
        ("使用步骤", "只在“数据”sheet 填写；“填写示例”不会被导入。上传后先核对最终企业树，再确认创建。"),
        ("必填字段", "客户名称、企业代码、项目简称、审计年度、项目类型、会计准则。"),
        ("企业代码", "18位统一社会信用代码；上级和最终控制方代码也使用同一规则。"),
        ("上级企业代码(parent)", "选填：填写本企业直接上级的统一社会信用代码。"),
        ("最终控制方代码(ultimate)", "选填：同一集团、同一年度填写相同的最终控制方代码。"),
        ("与上级关系(relation)", "有上级时选择子公司或分公司；留空时平台按企业名称判断。"),
        ("集团关系", "上级代码指向直接上级；最终控制方用于归组；关系只能是子公司或分公司。"),
        ("已有项目", "普通模板只新建。现有集团树请从合并模块导出后，将需要修改的行设为“更新集团关系”。"),
        ("数据上限", f"最多 {MAX_IMPORT_ROWS} 行；不得移动、重命名或删除表头。"),
    ]
    for item in notes_rows:
        notes.append(item)
    notes.column_dimensions["A"].width = 22
    notes.column_dimensions["B"].width = 110

    _add_options_and_validations(workbook, data)

    meta = workbook.create_sheet("_GT_META")
    meta.append(["template_version", TEMPLATE_VERSION])
    meta.append(["mode", mode])
    meta.sheet_state = "hidden"
    output = BytesIO()
    workbook.save(output)
    output.seek(0)
    return output

def parsed_to_excel_row(row: ParsedRow) -> list[object]:
    return [
        row.client_name, row.company_code, row.short_name, row.audit_year or "",
        PROJECT_TYPE_REVERSE.get(row.project_type or "", ""),
        ACCOUNTING_STANDARD_REVERSE.get(row.accounting_standard or "", ""),
        REPORT_SCOPE_REVERSE.get(row.report_scope, "单户"),
        row.parent_company_code, row.ultimate_company_code,
        RELATION_REVERSE.get(row.relation_to_parent or "", ""),
        ACTION_REVERSE.get(row.action, row.action),
        str(row.project_id) if row.project_id else "", row.base_hash,
    ]


def generate_error_workbook(
    rows: list[tuple[ParsedRow, list[str]]], *, global_errors: list[str] | None = None,
) -> BytesIO:
    workbook = Workbook()
    sheet = workbook.active
    sheet.title = "错误明细"
    sheet.append(EXTENDED_COLUMNS + ["错误原因"])
    for row, errors in rows:
        _append_text_safe(sheet, parsed_to_excel_row(row) + ["；".join(errors)])
    sheet.freeze_panes = "A2"
    sheet.auto_filter.ref = f"A1:N{max(sheet.max_row, 2)}"
    sheet.column_dimensions["N"].width = 70
    for cell in sheet[1]:
        cell.font = Font(bold=True)
        cell.fill = _HEADER_FILL
    for cell in sheet["N"][1:]:
        cell.fill = PatternFill("solid", fgColor="FCE4D6")
    if global_errors:
        notes = workbook.create_sheet("批次错误")
        notes.append(["错误原因"])
        for message in global_errors:
            notes.append([message])
        notes.column_dimensions["A"].width = 100
    output = BytesIO()
    workbook.save(output)
    output.seek(0)
    return output


def entity_group_fields_hash(projects: list[object]) -> str:
    rows = []
    for project in sorted(
        projects,
        key=lambda item: (
            str(getattr(item, "report_scope", None) or "standalone"),
            str(getattr(item, "id", "")),
        ),
    ):
        rows.append({
            "id": str(getattr(project, "id", "")),
            "scope": str(getattr(project, "report_scope", None) or "standalone"),
            "parent": str(getattr(project, "parent_company_code", None) or "").strip(),
            "ultimate": str(getattr(project, "ultimate_company_code", None) or "").strip(),
            "relation": str(getattr(project, "relation_to_parent", None) or "").strip(),
        })
    encoded = json.dumps(rows, ensure_ascii=False, sort_keys=True, separators=(",", ":")).encode("utf-8")
    return hashlib.sha256(encoded).hexdigest()


def group_fields_hash(project: object) -> str:
    payload = "\x1f".join(
        str(getattr(project, key, None) or "").strip()
        for key in ("company_code", "audit_year", "report_scope", "parent_company_code", "ultimate_company_code", "relation_to_parent")
    )
    return hashlib.sha256(payload.encode("utf-8")).hexdigest()


def export_workbook(rows: list[list], *, mode: str, root_project_id: str | None = None) -> BytesIO:
    workbook = Workbook()
    data = workbook.active
    data.title = "数据"
    data.append(EXTENDED_COLUMNS)
    for row in rows:
        _append_text_safe(data, row)
    _style_data_sheet(data)
    notes = workbook.create_sheet("说明事项")
    for item in (
        ("使用方式", "只修改需要调整的企业行；未变化的行保持“更新集团关系”，预校验会归为无变化。"),
        ("允许修改", "上级企业代码、最终控制方代码、与上级关系、操作类型；新增企业可填写完整新行。"),
        ("禁止修改", "已有行的客户名称、企业代码、简称、年度、项目类型、会计准则、报表类型。"),
        ("系统字段", "项目ID和关系基线摘要已隐藏并锁定；删除或修改后不能更新原项目。"),
        ("安全确认", "上传后必须核对最终企业树；确认执行前不会修改任何项目。"),
    ):
        notes.append(item)
    notes.column_dimensions["A"].width = 18
    notes.column_dimensions["B"].width = 110
    _add_options_and_validations(workbook, data)
    meta = workbook.create_sheet("_GT_META")
    meta.append(["template_version", TEMPLATE_VERSION])
    meta.append(["mode", mode])
    meta.append(["root_project_id", root_project_id or ""])
    meta.sheet_state = "hidden"
    output = BytesIO()
    workbook.save(output)
    output.seek(0)
    return output


__all__ = [
    "ACCOUNTING_STANDARD_MAP", "ACCOUNTING_STANDARD_REVERSE", "ACTION_MAP", "ACTION_REVERSE",
    "BASE_COLUMNS", "EXTENDED_COLUMNS", "LoadedWorkbook", "MAX_IMPORT_ROWS", "MAX_UPLOAD_BYTES",
    "ParsedRow", "PROJECT_TYPE_MAP", "PROJECT_TYPE_REVERSE", "RELATION_MAP", "RELATION_REVERSE",
    "REPORT_SCOPE_MAP", "REPORT_SCOPE_REVERSE", "TEMPLATE_VERSION", "export_workbook",
    "generate_error_workbook", "generate_template", "group_fields_hash", "entity_group_fields_hash",
    "is_blank_row", "load_data_sheet",
    "parse_row", "parsed_to_excel_row",
]

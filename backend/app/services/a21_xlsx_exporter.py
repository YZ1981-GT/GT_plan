"""A21~A25 复核表 xlsx 导出 — 模板回填 checklist_responses。"""

from __future__ import annotations

import io
import json
import logging
from pathlib import Path
from uuid import UUID

import openpyxl
import sqlalchemy as sa
from sqlalchemy.ext.asyncio import AsyncSession

from app.services.review_checklist_service import get_review_definition_for_wp, resolve_template_key

_logger = logging.getLogger(__name__)

_AUDIT_PATH = Path(__file__).resolve().parent.parent.parent / "data" / "a21_a25_xlsx_audit.json"
_TEMPLATES_DIR = Path(__file__).resolve().parent.parent.parent / "wp_templates" / "A"


def _load_audit_entries() -> list[dict]:
    if not _AUDIT_PATH.exists():
        return []
    return json.loads(_AUDIT_PATH.read_text(encoding="utf-8"))


def _find_audit_entry(wp_code: str, is_large_soe: bool) -> dict | None:
    key = resolve_template_key(wp_code, is_large_soe, "financial")
    for entry in _load_audit_entries():
        ek = entry.get("wp_code")
        variant = entry.get("enterprise_variant")
        if variant:
            ek = f"{ek}:{variant}"
        if ek == key or entry.get("wp_code") == wp_code:
            return entry
    return None


async def export_review_checklist_xlsx(
    db: AsyncSession,
    project_id: UUID,
    wp_id: UUID,
    wp_code: str,
) -> tuple[bytes, str]:
    """导出复核表 xlsx；返回 (bytes, filename)。"""
    from app.models.core import Project

    proj = await db.get(Project, project_id)
    is_soe = bool(getattr(proj, "is_large_soe", False)) if proj else False

    audit = _find_audit_entry(wp_code, is_soe)
    if not audit:
        raise FileNotFoundError(f"未找到 {wp_code} 的 audit 映射")

    filename = audit.get("filename") or f"{wp_code}.xlsx"
    template_path = _TEMPLATES_DIR / filename
    if not template_path.exists():
        raise FileNotFoundError(f"模板文件不存在: {filename}")

    r = await db.execute(
        sa.text(
            "SELECT item_id, conclusion, remark FROM checklist_responses WHERE wp_id = :wp_id"
        ),
        {"wp_id": str(wp_id)},
    )
    responses = {row.item_id: row for row in r.fetchall()}

    wb = openpyxl.load_workbook(template_path)
    checklist = audit.get("sheets", {}).get("checklist", {})
    sheet_name = checklist.get("name")
    ws = wb[sheet_name] if sheet_name and sheet_name in wb.sheetnames else wb.active

    cols = checklist.get("column_roles", {})
    col_yes = openpyxl.utils.column_index_from_string(cols.get("yes", "F"))
    col_no = openpyxl.utils.column_index_from_string(cols.get("no", "G"))
    col_na = openpyxl.utils.column_index_from_string(cols.get("na", "H"))
    col_ref = openpyxl.utils.column_index_from_string(cols.get("record_ref", "I"))

    for item in checklist.get("items", []):
        item_id = item.get("item_id")
        row = item.get("row")
        if not item_id or not row:
            continue
        resp = responses.get(item_id)
        if not resp:
            continue
        concl = (resp.conclusion or "").upper()
        ws.cell(row=row, column=col_yes, value="√" if concl == "Y" else None)
        ws.cell(row=row, column=col_no, value="√" if concl == "N" else None)
        ws.cell(row=row, column=col_na, value="N/A" if concl in ("NA", "N/A") else None)
        if resp.remark:
            ws.cell(row=row, column=col_ref, value=resp.remark)

    record_key = f"{wp_code}-record"
    if record_key in responses and responses[record_key].remark:
        record_sheet = audit.get("sheets", {}).get("record", {}).get("name")
        if record_sheet and record_sheet in wb.sheetnames:
            rs = wb[record_sheet]
            rs["A1"] = responses[record_key].remark

    sign_key = f"{wp_code}-sign"
    sign_resp = responses.get(sign_key)
    if sign_resp and checklist.get("sign_rows"):
        for sr in checklist["sign_rows"]:
            label = "通过" if sign_resp.conclusion == "pass" else "退回"
            ws.cell(row=sr, column=col_yes, value=label)

    buf = io.BytesIO()
    wb.save(buf)
    return buf.getvalue(), filename

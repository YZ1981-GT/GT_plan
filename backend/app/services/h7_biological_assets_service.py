"""H7 生产性生物资产 — 服务层.

Spec: .kiro/specs/h7-biological-assets/ Task 5.3
Requirements: 11.1-11.4, 13.1-13.3

业务逻辑：
- 行业适用性判断（agriculture/forestry/livestock/fishery）
- 导出模板/数据
- 导入数据
- 折旧验证（直线法）
- 互转验证（三方向差额为0）
"""
from __future__ import annotations

import io
import logging
from typing import Any
from urllib.parse import quote

import sqlalchemy as sa
from sqlalchemy.ext.asyncio import AsyncSession
from fastapi import UploadFile

logger = logging.getLogger(__name__)

# ─── Constants ────────────────────────────────────────────────────────────────

APPLICABLE_INDUSTRIES = {"agriculture", "forestry", "livestock", "fishery"}
ACCOUNT_CODE_1621 = "1621"


# ─── Industry Check ──────────────────────────────────────────────────────────

async def check_industry_applicability(wp_id: str, db: AsyncSession) -> dict[str, Any]:
    """检查项目行业适用性.

    Returns:
        {"is_applicable": bool, "industry": str, "message": str}
    """
    try:
        result = await db.execute(
            sa.text("""
                SELECT p.business_category
                FROM projects p
                JOIN wp_index wi ON wi.project_id = p.id
                JOIN working_paper wp ON wp.wp_index_id = wi.id
                WHERE wp.id = :wp_id
            """),
            {"wp_id": wp_id},
        )
        row = result.fetchone()
        if row:
            industry = (row.business_category or "").strip().lower()
            is_applicable = industry in APPLICABLE_INDUSTRIES
            return {
                "is_applicable": is_applicable,
                "industry": industry,
                "message": "" if is_applicable else "本底稿仅适用于农林牧渔行业项目",
            }
    except Exception as e:  # noqa: BLE001
        logger.warning("H7 industry check error: %s", e)

    return {"is_applicable": True, "industry": "", "message": ""}


# ─── Export Template ──────────────────────────────────────────────────────────

async def export_h7_template(
    wp_id: str,
    sheet: str | None,
    db: AsyncSession,
) -> tuple[io.BytesIO, str]:
    """导出H7空模板.

    TODO: 从 render_schema YAML 生成 xlsx 模板
    """
    buf = io.BytesIO()
    # Placeholder - 生成简单空 xlsx
    try:
        import openpyxl
        wb = openpyxl.Workbook()
        ws = wb.active
        ws.title = sheet or "H7-模板"
        ws.append(["编号", "项目", "金额"])
        wb.save(buf)
    except ImportError:
        buf.write(b"")
    buf.seek(0)
    filename = quote(f"H7生产性生物资产_模板_{sheet or 'all'}.xlsx")
    return buf, filename


# ─── Export Data ──────────────────────────────────────────────────────────────

async def export_h7_data(
    wp_id: str,
    sheet: str | None,
    db: AsyncSession,
) -> tuple[io.BytesIO, str]:
    """导出H7已填数据.

    TODO: 从 checklist_responses 提取数据写入 xlsx
    """
    buf = io.BytesIO()
    try:
        import openpyxl
        wb = openpyxl.Workbook()
        ws = wb.active
        ws.title = sheet or "H7-数据"
        ws.append(["编号", "项目", "金额", "备注"])

        # 读取 checklist_responses 中 H7 前缀的数据
        result = await db.execute(
            sa.text("""
                SELECT item_id, conclusion, remark
                FROM checklist_responses
                WHERE wp_id = :wp_id AND item_id LIKE 'H7-%'
                ORDER BY item_id
            """),
            {"wp_id": wp_id},
        )
        for row in result.fetchall():
            ws.append([row.item_id, "", row.conclusion or "", row.remark or ""])

        wb.save(buf)
    except ImportError:
        buf.write(b"")
    buf.seek(0)
    filename = quote(f"H7生产性生物资产_数据_{sheet or 'all'}.xlsx")
    return buf, filename


# ─── Import Data ──────────────────────────────────────────────────────────────

async def import_h7_data(
    wp_id: str,
    file: UploadFile,
    sheet: str | None,
    db: AsyncSession,
) -> dict[str, Any]:
    """导入H7数据.

    读取上传的 xlsx，解析并写入 checklist_responses。
    """
    content = await file.read()
    imported_count = 0

    try:
        import openpyxl
        wb = openpyxl.load_workbook(io.BytesIO(content), read_only=True)
        ws = wb.active
        for row_idx, row in enumerate(ws.iter_rows(min_row=2, values_only=True), start=2):
            if not row or not row[0]:
                continue
            item_id = str(row[0])
            if not item_id.startswith("H7-"):
                continue
            conclusion = str(row[2]) if len(row) > 2 and row[2] else None
            remark = str(row[3]) if len(row) > 3 and row[3] else None

            await db.execute(
                sa.text("""
                    INSERT INTO checklist_responses (wp_id, item_id, conclusion, remark)
                    VALUES (:wp_id, :item_id, :conclusion, :remark)
                    ON CONFLICT (wp_id, item_id)
                    DO UPDATE SET conclusion = :conclusion, remark = :remark
                """),
                {
                    "wp_id": wp_id,
                    "item_id": item_id,
                    "conclusion": conclusion,
                    "remark": remark,
                },
            )
            imported_count += 1
        await db.flush()
    except ImportError:
        logger.error("openpyxl not installed, cannot import")
    except Exception as e:  # noqa: BLE001
        logger.error("H7 import error: %s", e)

    return {"imported_count": imported_count, "sheet": sheet}


# ─── Depreciation Validation ─────────────────────────────────────────────────

def validate_straight_line_depreciation(
    cost: float,
    salvage_rate: float,
    useful_life: float,
    monthly_dep: float,
) -> dict[str, Any]:
    """验证直线法折旧计算正确性.

    Returns:
        {"is_valid": bool, "expected_monthly": float, "diff": float}
    """
    if useful_life <= 0:
        return {"is_valid": False, "expected_monthly": 0, "diff": abs(monthly_dep)}
    annual = cost * (1 - salvage_rate) / useful_life
    expected_monthly = annual / 12
    diff = abs(monthly_dep - expected_monthly)
    return {
        "is_valid": diff < 0.01,
        "expected_monthly": expected_monthly,
        "diff": diff,
    }


# ─── Transfer Validation ─────────────────────────────────────────────────────

def validate_transfer_balance(transfer_out: float, transfer_in: float) -> dict[str, Any]:
    """验证互转差额为0.

    Returns:
        {"is_valid": bool, "diff": float}
    """
    diff = transfer_out - transfer_in
    return {"is_valid": abs(diff) < 0.01, "diff": diff}

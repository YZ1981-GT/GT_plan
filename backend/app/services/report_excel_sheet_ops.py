"""财务报表 Excel 导出的工作表级辅助（spec ``deliverable-trio-generation-accuracy`` §1；由 ``ReportExcelExporter`` 拆出）。

- 公式格判定、金额取值（导出器各填充路径共用）；
- 模板自算行：分项全空时写库内同名行总额（国企「实收资本」）、单体口径清空「*」合并专用行公式（少数股东损益）；
- 表头写死年份按本年 / 上年改写；减值表期末列有值时写值（不留「年初 + 增加 − 减少」公式）；
- 附注列 ``{{note_ref:…}}`` 渲染与「有金额但附注无内容」提醒。

只操作已加载的工作簿与内存数据，不访问数据库。
"""

from __future__ import annotations

import re
from typing import Any

from openpyxl.cell.cell import MergedCell

__all__ = [
    "NOTE_REF_ANY_RE",
    "apply_impairment_closing",
    "blank_consolidated_only_formulas",
    "fill_formula_rows_without_inputs",
    "formula_ref_cells",
    "is_formula_cell",
    "normalize_header_years",
    "note_ref_warnings",
    "render_note_refs",
    "resolve_amount",
    "row_label_left_of",
]

#   合并模板附注列复合占位：{{note_ref:BS-002}};{{note_ref:BS-002:parent}}（合并附注 + 母公司附注）
NOTE_REF_ANY_RE = re.compile(r"\{\{note_ref:([^}:]+)(?::(parent))?\}\}")
_NOTE_REF_SEP_TRIM_RE = re.compile(r"^[\s;；,，、]+|[\s;；,，、]+$")


def is_formula_cell(cell) -> bool:
    """判断单元格是否为公式格（不可覆盖）."""
    if getattr(cell, "data_type", None) == "f":
        return True
    val = cell.value
    return isinstance(val, str) and val.startswith("=")


def resolve_amount(row: dict | None, period: str) -> float | None:
    """取 current/prior 金额（float 或 None）."""
    if not row:
        return None
    key = "current_period_amount" if period == "current" else "prior_period_amount"
    amount = row.get(key)
    if amount is None:
        return None
    try:
        return float(amount)
    except (TypeError, ValueError):
        return None


def formula_ref_cells(formula: str) -> list[tuple[int, int]] | None:
    """公式引用的本表单元格 [(行, 列)]；跨表 / 命名区域 / 解析失败返回 ``None``。"""
    from openpyxl.formula import Tokenizer
    from openpyxl.utils.cell import range_boundaries

    try:
        tokens = Tokenizer(formula).items
    except Exception:  # noqa: BLE001
        return None
    cells: list[tuple[int, int]] = []
    for t in tokens:
        if t.type != "OPERAND" or t.subtype != "RANGE":
            continue
        ref = t.value.replace("$", "")
        if "!" in ref:
            return None
        try:
            min_col, min_row, max_col, max_row = range_boundaries(ref)
        except (TypeError, ValueError):
            return None
        if None in (min_col, min_row, max_col, max_row):
            return None
        for r in range(min_row, max_row + 1):
            for c in range(min_col, max_col + 1):
                cells.append((r, c))
    return cells


#: 附注列已写入的编号（「七、58」「八、5;十二、1」），找行名时跳过
_NOTE_REF_TEXT_RE = re.compile(r"^(?:[一二三四五六七八九十百]+、\d+\s*[;；,，、]?\s*)+$")


def row_label_left_of(ws, row_no: int, col_no: int) -> str:
    """同行、指定列左侧最近的行名文字格（排除占位 / 公式 / 纯数字 / 已填的附注编号）。"""
    for col in range(col_no - 1, 0, -1):
        v = ws.cell(row_no, col).value
        if not isinstance(v, str):
            continue
        t = v.strip()
        if not t or "{{" in t or t.startswith("=") or re.fullmatch(r"[\d.\-—]+", t):
            continue
        if _NOTE_REF_TEXT_RE.fullmatch(t):
            continue
        return t
    return ""


#: 表头写死的年份：「2024年度」「2025年12月31日」
_HEADER_YEAR_RE = re.compile(r"^(\d{4})(年度|年12月31日)$")


def normalize_header_years(ws, year: int) -> None:
    """表头区（前 5 行）写死的年份按本年 / 上年替换。

    权益变动表上年块（列位于「上年金额 / 上期金额」表头之后）写上年，其余写本年。
    只改整格恰为「YYYY年度 / YYYY年12月31日」的文字格，不碰其它内容。
    """
    prior_col: int | None = None
    for row in ws.iter_rows(min_row=1, max_row=6):
        for c in row:
            if isinstance(c.value, str) and re.sub(r"\s", "", c.value) in ("上年金额", "上期金额"):
                prior_col = c.column if prior_col is None else min(prior_col, c.column)
    for row in ws.iter_rows(min_row=1, max_row=5):
        for c in row:
            if isinstance(c, MergedCell) or not isinstance(c.value, str):
                continue
            m = _HEADER_YEAR_RE.match(c.value.strip())
            if not m:
                continue
            target = year - 1 if (prior_col is not None and c.column >= prior_col) else year
            if int(m.group(1)) != target:
                c.value = f"{target}{m.group(2)}"


def fill_formula_rows_without_inputs(ws, best_by_key: dict, family: str) -> int:
    """模板自算行的分项全空、库内有同名行金额时，写库内金额（覆盖公式）。

    国企资产负债表「实收资本」= SUM(国家资本…外商资本)，库内只有实收资本总额、没有构成
    明细 ⇒ 模板公式算出 0，实收资本 1.5 亿整行丢失、负债和所有者权益总计随之少 1.5 亿
    （2026-10-16 重药安徽实测）。分项有任一非零时保留模板公式（按明细自算）。
    """
    if not best_by_key or family not in ("BS", "IS", "CFS"):
        return 0
    from app.services.report_template_binding import normalize_row_label, period_columns

    memo: dict[tuple[int, int], bool] = {}

    def _empty(r: int, c: int, depth: int = 0) -> bool:
        """单元格为空 / 0，或为本表公式且其全部输入（递归）为空。"""
        if (r, c) in memo:
            return memo[(r, c)]
        memo[(r, c)] = False  # 环引用保护：在途按非空处理
        cell = ws.cell(r, c)
        v = cell.value
        if is_formula_cell(cell):
            refs = formula_ref_cells(str(v)) if depth < 20 else None
            result = bool(refs) and all(_empty(rr, cc, depth + 1) for rr, cc in refs if (rr, cc) != (r, c))
        else:
            result = v in (None, "") or (isinstance(v, (int, float)) and v == 0)
        memo[(r, c)] = result
        return result

    cur_col, prior_col = period_columns(ws)
    written = 0
    for col, period in ((cur_col, "current"), (prior_col, "prior")):
        if col is None:
            continue
        for row_no in range(1, ws.max_row + 1):
            cell = ws.cell(row_no, col)
            if isinstance(cell, MergedCell) or not is_formula_cell(cell):
                continue
            refs = formula_ref_cells(str(cell.value))
            if not refs:
                continue
            if not all(_empty(r, c) for r, c in refs if (r, c) != (row_no, col)):
                continue
            label = row_label_left_of(ws, row_no, col)
            key, markers = normalize_row_label(label)
            src = best_by_key.get((family, key, markers)) if key else None
            amount = resolve_amount(src, period)
            if amount is None or amount == 0:
                continue
            cell.value = amount
            memo[(row_no, col)] = False
            written += 1
    return written


def blank_consolidated_only_formulas(ws) -> int:
    """单体口径：模板注明「带*科目为合并会计报表专用」，这些行的模板自算公式清空。

    国企利润表「*2.少数股东损益 = 净利润 - 归属于母公司所有者的净利润」：单体没有归属于
    母公司的拆分，公式会把全部净利润算成少数股东损益（2026-10-16 四个单体项目实测）。
    只清公式格；库内填了数的格原样保留。
    """
    from app.services.report_template_binding import normalize_row_label, period_columns

    cols = [c for c in period_columns(ws) if c is not None]
    if not cols:
        return 0
    cleared = 0
    for row_no in range(1, ws.max_row + 1):
        label = row_label_left_of(ws, row_no, min(cols))
        if not label or "*" not in normalize_row_label(label)[1]:
            continue
        for col in cols:
            cell = ws.cell(row_no, col)
            if not isinstance(cell, MergedCell) and is_formula_cell(cell):
                cell.value = None
                cleared += 1
    return cleared


def apply_impairment_closing(ws, binding: Any, parent_binding: Any) -> None:
    """减值表期末列：模板为「=年初+增加-减少」公式，而本期增减通常无数据源 ⇒

    公式会把期末算成年初数（错数）。有期末值时直接写值，增减列保持空白待补。
    """
    if binding is None:
        return
    closing_col: int | None = None
    for row in ws.iter_rows(min_row=1, max_row=6):
        for c in row:
            if isinstance(c.value, str) and "期末" in c.value:
                closing_col = c.column
    if closing_col is None:
        return
    src = parent_binding if ("母公司" in ws.title and parent_binding is not None) else binding
    imp_re = re.compile(r"\{\{imp:([^:}]+):")
    for slot in [s for s in getattr(src, "by_slot", {})]:
        parts = slot.split("!")
        if len(parts) != 3 or parts[0] != ws.title or not parts[2].startswith("IMP-"):
            continue
        row_idx = int(parts[1])
        row = src.by_slot[slot]
        amount = resolve_amount(row, "current")
        if amount is None:
            continue
        cell = ws.cell(row_idx, closing_col)
        if isinstance(cell, MergedCell):
            continue
        if is_formula_cell(cell) or (isinstance(cell.value, str) and imp_re.search(cell.value)) or cell.value is None:
            cell.value = amount


def note_ref_warnings(
    note_plan: Any,
    misses: list[tuple[str, int, str, str]],
    binding: Any,
    report_data: dict[str, list[dict]],
) -> list[str]:
    """附注列提醒：附注未生成 / 报表有金额但附注对应章节无内容。"""
    if note_plan is None:
        return []
    if not note_plan.has_notes:
        return ["附注尚未生成：报表「附注」列未填写编号，生成附注后请重新生成财务报表。"]
    row_index = {
        r.get("row_code"): r for rows in report_data.values() for r in rows if r.get("row_code")
    }
    labels: list[str] = []
    for sheet, row_no, code, label in misses:
        src = binding.lookup(sheet, row_no, code) if binding is not None else row_index.get(code)
        if not src:
            continue
        amounts = [resolve_amount(src, "current"), resolve_amount(src, "prior")]
        if not any(a for a in amounts if a is not None):
            continue
        name = re.sub(r"\s+", "", label or code)
        if name not in labels:
            labels.append(name)
    if not labels:
        return []
    shown = "、".join(labels[:12]) + ("等" if len(labels) > 12 else "")
    return [f"附注列：{len(labels)} 个报表项目有金额，但附注对应章节无内容，附注列留空：{shown}。"]


def render_note_refs(text: str, label: str, note_refs: dict[str, str], note_plan: Any) -> tuple[str | None, list[str]]:
    """替换格内全部 ``{{note_ref:…}}``；取不到编号的占位写空，并清掉多余分隔符。

    ``label`` 为该格左侧行名（按行名找附注节）。返回 ``(格值, 取不到编号的合并附注占位编码)``。
    """
    missed: list[str] = []

    def _one(m: re.Match) -> str:
        code, is_parent = m.group(1), bool(m.group(2))
        if note_plan is None:
            return "" if is_parent else (note_refs.get(code) or "")
        ref = note_plan.ref_for(label, parent=is_parent, hint_code=code)
        if ref is None and not is_parent:
            missed.append(code)
        return ref or ""

    out = _NOTE_REF_SEP_TRIM_RE.sub("", NOTE_REF_ANY_RE.sub(_one, text))
    out = re.sub(r"[;；]{2,}", ";", out)
    return out or None, missed

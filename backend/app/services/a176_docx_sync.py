"""A17-6 总结会会议纪要 — 结构化数据 ↔ Word 双向同步服务

实物模板结构（以 Table[1] 为准）:
- Table[0]: 提示横幅（只读，不同步）
- Table[1] 11×3:
  R0  被审计单位名称 / 索引号
  R1  报表截止日/期间 / 编制人 / 日期
  R2  内控评价基准日 / 复核人 / 日期
  R4  会议地点
  R5  会议时间
  R6  会议组织者
  R7  出席会议者（→ attendees）
  R8  记录员
  R9  会议议题标题行
  R10 全部议程正文（编号 1、…10、 塞在同一单元格）

设计原则: 结构化视图权威；OO 往返只读写 Table[1]，不再依赖段落编号。
"""

from __future__ import annotations

import logging
import re
import shutil
from pathlib import Path
from typing import Optional
from uuid import UUID

import sqlalchemy as sa
from sqlalchemy.ext.asyncio import AsyncSession

logger = logging.getLogger(__name__)

_BACKEND_ROOT = Path(__file__).resolve().parent.parent.parent
TEMPLATE_PATH = _BACKEND_ROOT / "wp_templates" / "A" / "A17-6  总结会会议记要.docx"
STORAGE_BASE = _BACKEND_ROOT / "storage" / "projects"

# UI 议题标题（写入 docx 时用）
AGENDA_TITLES: dict[int, str] = {
    1: "总述审计过程",
    2: "针对计划阶段评估的重大错报风险采取的应对措施及审计结论",
    3: "针对评估的特别风险采取的应对措施及审计结论",
    4: "内部控制测试和审计风险",
    5: "已发现的错报和重要情况汇总",
    6: "其他重要的会计和审计事项",
    7: "会计师事务所资源与独立性",
    8: "确定拟发表的审计意见",
    9: "拟发表的审计意见及关键审计事项",
    10: "其他",
}

_AGENDA_SPLIT = re.compile(r"(?=(?:^|\n)\s*(\d{1,2})[、.．]\s*)")


def _get_project_file(project_id: UUID) -> Path:
    return STORAGE_BASE / str(project_id) / "workpapers" / "A" / "A17-6.docx"


def _ensure_project_file(project_id: UUID) -> Path:
    fp = _get_project_file(project_id)
    if not fp.exists():
        fp.parent.mkdir(parents=True, exist_ok=True)
        if not TEMPLATE_PATH.exists():
            raise FileNotFoundError(f"A17-6 模板不存在: {TEMPLATE_PATH}")
        shutil.copy2(TEMPLATE_PATH, fp)
    return fp


def _meta_table(doc):
    """返回业务元信息表（Table[1]；若仅一表则回退 Table[0]）。"""
    if len(doc.tables) >= 2:
        return doc.tables[1]
    return doc.tables[0] if doc.tables else None


def _cell_text(row, col: int) -> str:
    if col >= len(row.cells):
        return ""
    return (row.cells[col].text or "").strip()


def _set_cell(row, col: int, value: str) -> None:
    if col >= len(row.cells):
        return
    row.cells[col].text = value or ""


def _parse_agenda_blob(blob: str) -> dict[int, str]:
    """将 R10 单元格中的「1、…\\n2、…」拆成 {1: content, ...}。"""
    text = (blob or "").replace("\r\n", "\n").strip()
    if not text:
        return {}

    # 在编号前切开
    parts = re.split(r"(?=\d{1,2}[、.．])", text)
    result: dict[int, str] = {}
    for part in parts:
        part = part.strip()
        if not part:
            continue
        m = re.match(r"^(\d{1,2})[、.．]\s*(.*)$", part, re.DOTALL)
        if not m:
            continue
        idx = int(m.group(1))
        if not (1 <= idx <= 10):
            continue
        body = (m.group(2) or "").strip()
        # 去掉标题行（若正文以标题开头）
        title = AGENDA_TITLES.get(idx, "")
        if title and body.startswith(title):
            body = body[len(title):].lstrip("：: \n")
        result[idx] = body
    return result


def _format_agenda_blob(agenda: dict[int, str]) -> str:
    """将议程 dict 写回单一单元格文本。"""
    lines: list[str] = []
    for i in range(1, 11):
        title = AGENDA_TITLES.get(i, f"议题{i}")
        content = (agenda.get(i) or "").strip()
        if content:
            lines.append(f"{i}、{title}\n{content}")
        else:
            lines.append(f"{i}、{title}")
    return "\n\n".join(lines)


def parse_docx(file_path: Path) -> tuple[dict[int, str], dict[str, str]]:
    """从 docx Table[1] 解析议程 + 元信息。"""
    from docx import Document

    doc = Document(str(file_path))
    table = _meta_table(doc)
    meta: dict[str, str] = {}
    agenda: dict[int, str] = {}
    if table is None:
        return agenda, meta

    rows = table.rows
    n = len(rows)

    def row_blob(ri: int) -> str:
        if ri >= n:
            return ""
        # 合并单元格时多列文本相同，取第一格即可
        return _cell_text(rows[ri], 0)

    # R0: 被审计单位 / 索引
    if n > 0:
        r0 = rows[0]
        c0 = _cell_text(r0, 0)
        c2 = _cell_text(r0, 2) if len(r0.cells) > 2 else ""
        # 「被审计单位名称：XXX」或值在邻格
        if "：" in c0 or ":" in c0:
            meta["client_name"] = re.split(r"[：:]", c0, 1)[-1].strip()
        else:
            meta["client_name"] = _cell_text(r0, 1)
        if "索引" in c2:
            meta["index_no"] = re.split(r"[：:]", c2, 1)[-1].strip() or "A17-6"
        elif c2:
            meta["index_no"] = c2

    # R1: 期间 / 编制人 / 日期
    if n > 1:
        r1 = rows[1]
        period_cell = _cell_text(r1, 0)
        meta["period"] = re.split(r"[：:]", period_cell, 1)[-1].strip() if period_cell else ""
        prep = _cell_text(r1, 1)
        meta["preparer"] = re.split(r"[：:]", prep, 1)[-1].strip() if prep else ""
        meta["preparer_date"] = _cell_text(r1, 2).replace("日期：", "").replace("日期:", "").strip()

    # R2: 复核人
    if n > 2:
        r2 = rows[2]
        rev = _cell_text(r2, 1)
        meta["reviewer"] = re.split(r"[：:]", rev, 1)[-1].strip() if rev else ""
        meta["reviewer_date"] = _cell_text(r2, 2).replace("日期：", "").replace("日期:", "").strip()

    # R4–R8 会议字段（标签与值常在合并格）
    field_map = [
        (4, "meeting_place", ("会议地点",)),
        (5, "meeting_time", ("会议时间",)),
        (6, "organizer", ("会议组织者", "组织者")),
        (7, "attendees", ("出席会议者", "出席", "参会")),
        (8, "recorder", ("记录员", "记录")),
    ]
    for ri, key, labels in field_map:
        if ri >= n:
            continue
        raw = row_blob(ri)
        val = raw
        for lb in labels:
            if lb in raw:
                val = re.split(r"[：:]", raw, 1)[-1].strip()
                # 若分割后仍是标签本身，取第 2 列
                if not val or val == lb:
                    val = _cell_text(rows[ri], 1)
                    val = re.split(r"[：:]", val, 1)[-1].strip()
                break
        else:
            val = _cell_text(rows[ri], 1) or raw
        meta[key] = val

    # R10: 议程 blob
    if n > 10:
        agenda = _parse_agenda_blob(_cell_text(rows[10], 0))

    # 清理空 meta
    meta = {k: v for k, v in meta.items() if v}
    return agenda, meta


async def generate_docx(wp_id: UUID, project_id: UUID, db: AsyncSession) -> Path:
    """从 checklist_responses 填充 Table[1]，返回文件路径。"""
    from docx import Document

    file_path = _ensure_project_file(project_id)

    result = await db.execute(
        sa.text(
            "SELECT item_id, conclusion, remark FROM checklist_responses "
            "WHERE wp_id = :wp_id AND item_id LIKE 'a176-%'"
        ),
        {"wp_id": str(wp_id)},
    )

    meta_values: dict[str, str] = {}
    agenda_values: dict[int, str] = {}
    for row in result.fetchall():
        item_id = row.item_id
        if item_id.startswith("a176-meta-"):
            meta_values[item_id.removeprefix("a176-meta-")] = row.conclusion or row.remark or ""
        elif item_id.startswith("a176-agenda-"):
            try:
                idx = int(item_id.removeprefix("a176-agenda-"))
                if 1 <= idx <= 10:
                    agenda_values[idx] = row.remark or row.conclusion or ""
            except ValueError:
                pass

    doc = Document(str(file_path))
    table = _meta_table(doc)
    if table is None:
        doc.save(str(file_path))
        return file_path

    rows = table.rows
    n = len(rows)

    def write_labeled(ri: int, label: str, value: str, value_col: int = 0) -> None:
        if ri >= n or not value:
            return
        # 标签+值写回同一合并格，保持模板样式
        _set_cell(rows[ri], value_col, f"{label}{value}")
        # 同步合并副本列
        if len(rows[ri].cells) > 1 and value_col == 0:
            _set_cell(rows[ri], 1, f"{label}{value}")

    if n > 0 and meta_values.get("client_name"):
        write_labeled(0, "被审计单位名称：", meta_values["client_name"])
        if len(rows[0].cells) > 2:
            _set_cell(rows[0], 2, f"索引号：{meta_values.get('index_no') or 'A17-6'}")

    if n > 1:
        if meta_values.get("period"):
            write_labeled(1, "报表截止日/会计期间：", meta_values["period"])
        if meta_values.get("preparer"):
            _set_cell(rows[1], 1, f"编制人：{meta_values['preparer']}")
        if meta_values.get("preparer_date"):
            _set_cell(rows[1], 2, f"日期：{meta_values['preparer_date']}")

    if n > 2:
        if meta_values.get("reviewer"):
            _set_cell(rows[2], 1, f"复核人：{meta_values['reviewer']}")
        if meta_values.get("reviewer_date"):
            _set_cell(rows[2], 2, f"日期：{meta_values['reviewer_date']}")

    write_labeled(4, "会议地点：", meta_values.get("meeting_place", ""))
    write_labeled(5, "会议时间：", meta_values.get("meeting_time", ""))
    write_labeled(6, "会议组织者：", meta_values.get("organizer", ""))
    # 出席 ← attendees；兼容旧 convener
    attendees = meta_values.get("attendees") or meta_values.get("convener") or ""
    write_labeled(7, "出席会议者：", attendees)
    write_labeled(8, "记录员：", meta_values.get("recorder", ""))

    if n > 10:
        blob = _format_agenda_blob(agenda_values)
        _set_cell(rows[10], 0, blob)
        if len(rows[10].cells) > 1:
            _set_cell(rows[10], 1, blob)
        if len(rows[10].cells) > 2:
            _set_cell(rows[10], 2, blob)

    doc.save(str(file_path))
    logger.info("A17-6 docx generated (Table[1]): %s", file_path)
    return file_path


async def sync_docx_to_responses(
    wp_id: UUID, project_id: UUID, db: AsyncSession, user_id: UUID
) -> int:
    """从项目 docx Table[1] 解析并回写 checklist_responses。"""
    from datetime import datetime, timezone

    file_path = _get_project_file(project_id)
    if not file_path.exists():
        logger.warning("A17-6 docx 不存在，无法同步: %s", file_path)
        return 0

    agenda_dict, meta_dict = parse_docx(file_path)
    if not agenda_dict and not meta_dict:
        return 0

    now = datetime.now(timezone.utc)
    count = 0

    for key, value in meta_dict.items():
        if not value:
            continue
        item_id = f"a176-meta-{key}"
        await db.execute(
            sa.text("""
                INSERT INTO checklist_responses
                    (project_id, wp_id, item_id, conclusion, remark, updated_by, created_at, updated_at)
                VALUES (:pid, :wp_id, :item_id, :conclusion, NULL, :uid, :now, :now)
                ON CONFLICT (wp_id, item_id) DO UPDATE SET
                    conclusion = EXCLUDED.conclusion,
                    updated_by = EXCLUDED.updated_by,
                    updated_at = EXCLUDED.updated_at
            """),
            {
                "pid": str(project_id),
                "wp_id": str(wp_id),
                "item_id": item_id,
                "conclusion": value,
                "uid": str(user_id),
                "now": now,
            },
        )
        count += 1

    for idx, content in agenda_dict.items():
        if not content.strip():
            continue
        item_id = f"a176-agenda-{idx}"
        await db.execute(
            sa.text("""
                INSERT INTO checklist_responses
                    (project_id, wp_id, item_id, conclusion, remark, updated_by, created_at, updated_at)
                VALUES (:pid, :wp_id, :item_id, NULL, :remark, :uid, :now, :now)
                ON CONFLICT (wp_id, item_id) DO UPDATE SET
                    remark = EXCLUDED.remark,
                    updated_by = EXCLUDED.updated_by,
                    updated_at = EXCLUDED.updated_at
            """),
            {
                "pid": str(project_id),
                "wp_id": str(wp_id),
                "item_id": item_id,
                "remark": content,
                "uid": str(user_id),
                "now": now,
            },
        )
        count += 1

    await db.commit()
    logger.info("A17-6 docx→responses synced %d items (Table[1])", count)
    return count

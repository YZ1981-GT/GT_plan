"""N4 税金及附加 — 业务逻辑服务

损益类取数（tb_ledger本期发生额，与H10/I6/L8/N5同款）+ 多税种测算 + N2计提对应 + 跨底稿合计。

提供功能：
- 损益类P&L发生额（从tb_ledger取借方发生-贷方发生，非余额！）
- 多税种测算引擎（城建税及附加/房产税/印花税/土地使用税，与N2同源）
- N2应交税费计提额交叉验证（费用确认=计提额）
- A类利润表勾稽（审定发生额总计）
- 导出模板 / 导出数据 / 导入数据
- 审定数回写 trial_balance

科目6403税金及附加（**损益类/借方**）：本期发生额=借方发生-贷方发生
N税费循环损益类底稿（9 sheet/~110+公式，其中O2A原底稿标记skip）。

Requirements: 4.1-4.5, 6.1-6.5, 7.1-7.4
"""

from __future__ import annotations

import io
import json
import logging
from typing import Any
from uuid import uuid4

from openpyxl import Workbook, load_workbook
from openpyxl.styles import Alignment, Font, PatternFill

logger = logging.getLogger(__name__)

_N4_ACCOUNT_CODE = "6403"


# ═══════════════════════════════════════════════════════════════════════════════
# 多税种测算引擎（纯函数，与N2同源）
# ═══════════════════════════════════════════════════════════════════════════════


def calculate_surtax(vat: float, consumption_tax: float, rate: float) -> float:
    """城建税及附加 = (增值税 + 消费税) × 税率

    rate:
    - 城建税：市区7%(0.07) / 县城5%(0.05) / 其他1%(0.01)
    - 教育费附加：3%(0.03)
    - 地方教育附加：2%(0.02)

    Requirements: 4.1
    """
    return (vat + consumption_tax) * rate


def calculate_property_tax_by_value(original_value: float, deduct_rate: float) -> float:
    """房产税从价 = 房产原值 × (1 - 扣除比例) × 1.2%

    deduct_rate 通常为 10%~30% (0.1~0.3)。

    Requirements: 4.2
    """
    return original_value * (1 - deduct_rate) * 0.012


def calculate_property_tax_by_rent(rent_income: float) -> float:
    """房产税从租 = 租金收入 × 12%

    Requirements: 4.2
    """
    return rent_income * 0.12


def calculate_stamp_tax(taxable_amount: float, rate: float) -> float:
    """印花税 = 计税金额 × 适用税率

    按合同类型不同税率：
    - 购销合同：0.3‰ (0.0003)
    - 租赁合同：1‰ (0.001)
    - 借款合同：0.05‰ (0.00005)
    - 财产保险：1‰ (0.001)
    - 技术合同：0.3‰ (0.0003)

    Requirements: 4.3
    """
    return taxable_amount * rate


def calculate_land_use_tax(area: float, unit_tax: float) -> float:
    """土地使用税 = 占地面积 × 单位税额

    单位税额由各地自行确定（元/㎡·年），一般0.6~30元。

    Requirements: 4.4
    """
    return area * unit_tax


# ═══════════════════════════════════════════════════════════════════════════════
# 损益类公式引擎（纯函数）
# ═══════════════════════════════════════════════════════════════════════════════


def calc_period_amount(debit_occur: float, credit_occur: float) -> float:
    """损益类本期发生额 = 借方发生 - 贷方发生（6403费用类借方科目！）

    6403税金及附加为借方科目：借方=费用增加，贷方=费用冲回/红冲。
    净发生额=借方-贷方（正数=净费用），与H10/I6/L8/N5同款。

    Requirements: 7.1, 7.4
    """
    return debit_occur - credit_occur


def calc_audited_amount(unadjusted: float, aje: float, rje: float) -> float:
    """审定数 = 未审数 + AJE + RJE

    Requirements: 7.3
    """
    return unadjusted + aje + rje


def calc_yoy_change(current: float, prior: float) -> float | None:
    """同比变动 = (本期 - 上期) / 上期

    上期为0时返回 None（避免除零，前端显示"—"）。

    Requirements: 6.5
    """
    if prior == 0:
        return None
    return (current - prior) / prior


def calc_subtotal(amounts: list[float]) -> float:
    """合计行 = Σ各项金额

    Requirements: 6.3
    """
    return sum(amounts)


# ═══════════════════════════════════════════════════════════════════════════════
# 损益类取数（DB依赖）— 从 tb_ledger 取本期发生额
# ═══════════════════════════════════════════════════════════════════════════════


async def get_period_amounts(
    db: Any,
    project_id: Any,
    year: int,
) -> dict[str, Any]:
    """从 tb_ledger 取科目6403税金及附加本期发生额（损益类借方！）

    损益类科目取发生额（非余额！），从序时账按明细科目分组聚合。
    使用 account_code LIKE '6403%' 包含所有明细科目(640301消费税/640302城建税等)。
    返回总额及各明细科目(6403xx)的借方/贷方发生额。

    与H10资产处置损益、I6研发费用、L8财务费用、N5所得税费用同款取数逻辑。

    Requirements: 7.1, 7.2
    """
    import sqlalchemy as sa

    from app.models.audit_platform_models import TbLedger
    from app.services.dataset_query import get_active_filter

    result: dict[str, Any] = {
        "account_code": _N4_ACCOUNT_CODE,
        "account_name": "税金及附加",
        "direction": "debit",
        "total_debit": 0.0,
        "total_credit": 0.0,
        "net_amount": 0.0,
        "details": [],
        "found": False,
    }

    try:
        active_filter = await get_active_filter(
            db, TbLedger.__table__, project_id, year
        )

        # 1. 汇总：6403整体发生额
        stmt_total = sa.select(
            sa.func.coalesce(sa.func.sum(TbLedger.debit_amount), 0).label("total_debit"),
            sa.func.coalesce(sa.func.sum(TbLedger.credit_amount), 0).label("total_credit"),
        ).where(
            active_filter,
            TbLedger.account_code.startswith(_N4_ACCOUNT_CODE),
        )
        row = (await db.execute(stmt_total)).fetchone()
        if row:
            debit = _parse_num(row.total_debit)
            credit = _parse_num(row.total_credit)
            result["total_debit"] = round(debit, 2)
            result["total_credit"] = round(credit, 2)
            # 6403借方科目：净发生额=借方-贷方
            result["net_amount"] = round(debit - credit, 2)
            result["found"] = debit > 0 or credit > 0

        # 2. 明细：按子科目分组
        stmt_detail = (
            sa.select(
                TbLedger.account_code,
                sa.func.coalesce(sa.func.sum(TbLedger.debit_amount), 0).label("debit_total"),
                sa.func.coalesce(sa.func.sum(TbLedger.credit_amount), 0).label("credit_total"),
            )
            .where(
                active_filter,
                TbLedger.account_code.startswith(_N4_ACCOUNT_CODE),
            )
            .group_by(TbLedger.account_code)
            .order_by(TbLedger.account_code)
        )
        detail_rows = (await db.execute(stmt_detail)).fetchall()
        for d_row in detail_rows:
            d_debit = _parse_num(d_row.debit_total)
            d_credit = _parse_num(d_row.credit_total)
            result["details"].append({
                "account_code": d_row.account_code,
                "debit_amount": round(d_debit, 2),
                "credit_amount": round(d_credit, 2),
                "net_amount": round(d_debit - d_credit, 2),
            })

    except Exception as e:  # noqa: BLE001
        logger.warning(
            "N4 service: tb_ledger 损益类发生额查询失败 project_id=%s: %s",
            project_id, e,
        )

    return result


async def get_prior_year_amount(
    db: Any,
    project_id: Any,
    year: int,
) -> float:
    """从 trial_balance 取科目6403上期审定数（用于同比变动计算）.

    Requirements: 7.2
    """
    import sqlalchemy as sa

    prior_year = year - 1
    if prior_year <= 0:
        return 0.0

    try:
        result = await db.execute(
            sa.text("""
                SELECT COALESCE(SUM(audited_amount), 0) AS prior_audited
                FROM trial_balance
                WHERE project_id = :pid AND year = :year AND is_deleted = false
                  AND standard_account_code LIKE :code_prefix
            """),
            {
                "pid": str(project_id),
                "year": prior_year,
                "code_prefix": f"{_N4_ACCOUNT_CODE}%",
            },
        )
        row = result.fetchone()
        if row and row.prior_audited is not None:
            return _parse_num(row.prior_audited)
    except Exception as e:  # noqa: BLE001
        logger.warning("N4 service: 上期数查询失败: %s", e)

    return 0.0


# ═══════════════════════════════════════════════════════════════════════════════
# TB 回写
# ═══════════════════════════════════════════════════════════════════════════════


async def writeback_tb(
    db: Any,
    project_id: Any,
    audited_amount: float,
) -> dict[str, Any]:
    """回写审定数到 trial_balance（科目6403，本期发生额，损益类借方）.

    Requirements: 7.3
    """
    import sqlalchemy as sa

    try:
        result = await db.execute(
            sa.text(
                "UPDATE trial_balance "
                "SET audited_amount = :amount "
                "WHERE project_id = :pid AND standard_account_code = :code"
            ),
            {
                "amount": audited_amount,
                "pid": str(project_id),
                "code": _N4_ACCOUNT_CODE,
            },
        )
        await db.flush()
        return {
            "success": True,
            "account_code": _N4_ACCOUNT_CODE,
            "audited_amount": audited_amount,
            "rows_affected": result.rowcount,
        }
    except Exception as e:  # noqa: BLE001
        logger.error("N4 service: TB 回写失败 project_id=%s: %s", project_id, e)
        return {"success": False, "error": str(e)}


# ═══════════════════════════════════════════════════════════════════════════════
# N2 计提对应交叉验证
# ═══════════════════════════════════════════════════════════════════════════════


async def get_n2_accruals(
    db: Any,
    project_id: Any,
    year: int,
) -> dict[str, float]:
    """从 N2 底稿 checklist_responses 获取各税种本期计提额.

    查找 item_id 格式: N2-1-tax-accrual-{taxType}
    N2各税种计提额是N4费用确认的对应方（费用确认=计提额）。

    Requirements: 6.1, 6.3
    """
    import sqlalchemy as sa

    tax_accruals: dict[str, float] = {
        "消费税": 0.0,
        "城建税": 0.0,
        "教育费附加": 0.0,
        "地方教育附加": 0.0,
        "房产税": 0.0,
        "土地使用税": 0.0,
        "车船税": 0.0,
        "印花税": 0.0,
        "资源税": 0.0,
    }

    try:
        rows = (
            await db.execute(
                sa.text(
                    "SELECT cr.item_id, cr.conclusion "
                    "FROM checklist_responses cr "
                    "JOIN working_papers wp ON wp.id = cr.wp_id "
                    "WHERE wp.project_id = :pid "
                    "AND cr.item_id LIKE 'N2-1-tax-accrual-%'"
                ),
                {"pid": str(project_id)},
            )
        ).fetchall()

        for row in rows:
            item_id = row.item_id
            tax_type = item_id.replace("N2-1-tax-accrual-", "")
            amount = _parse_num(row.conclusion)
            if tax_type in tax_accruals:
                tax_accruals[tax_type] = amount

    except Exception as e:  # noqa: BLE001
        logger.warning("N4 service: N2计提额查询失败 project_id=%s: %s", project_id, e)

    return tax_accruals


def verify_expense_vs_accrual(
    n4_expenses: dict[str, float],
    n2_accruals: dict[str, float],
) -> list[dict[str, Any]]:
    """交叉验证：N4费用确认 vs N2计提额，逐税种对比返回差异.

    费用确认（N4）应等于本期计提额（N2）。差异>0.01时标记告警。

    Requirements: 6.1, 6.3, 6.5
    """
    results: list[dict[str, Any]] = []

    all_taxes = set(list(n4_expenses.keys()) + list(n2_accruals.keys()))
    for tax in sorted(all_taxes):
        expense = n4_expenses.get(tax, 0.0)
        accrual = n2_accruals.get(tax, 0.0)
        diff = expense - accrual
        results.append({
            "tax_type": tax,
            "n4_expense": round(expense, 2),
            "n2_accrual": round(accrual, 2),
            "difference": round(diff, 2),
            "is_matched": abs(diff) <= 0.01,
        })

    return results


# ═══════════════════════════════════════════════════════════════════════════════
# 跨底稿合计（A利润表勾稽）
# ═══════════════════════════════════════════════════════════════════════════════


async def get_income_statement_amount(
    db: Any,
    project_id: Any,
    year: int,
) -> dict[str, Any]:
    """获取N4审定发生额总计，供A类利润表税金及附加行勾稽.

    从 trial_balance 取科目6403合计审定额。
    若 trial_balance 尚未回写，则从 tb_ledger 实时聚合。

    Requirements: 6.4
    """
    import sqlalchemy as sa

    # 优先从 trial_balance 取已审定值
    try:
        result = await db.execute(
            sa.text("""
                SELECT COALESCE(SUM(audited_amount), 0) AS audited_total
                FROM trial_balance
                WHERE project_id = :pid AND year = :year AND is_deleted = false
                  AND standard_account_code LIKE :code_prefix
            """),
            {
                "pid": str(project_id),
                "year": year,
                "code_prefix": f"{_N4_ACCOUNT_CODE}%",
            },
        )
        row = result.fetchone()
        audited = _parse_num(row.audited_total) if row else 0.0

        if audited != 0:
            return {
                "account_code": _N4_ACCOUNT_CODE,
                "amount": round(audited, 2),
                "source": "trial_balance",
                "year": year,
            }
    except Exception as e:  # noqa: BLE001
        logger.warning("N4 service: trial_balance 取数失败: %s", e)

    # 回退：从 tb_ledger 实时聚合
    period_data = await get_period_amounts(db, project_id, year)
    return {
        "account_code": _N4_ACCOUNT_CODE,
        "amount": period_data["net_amount"],
        "source": "tb_ledger",
        "year": year,
    }


# ═══════════════════════════════════════════════════════════════════════════════
# 导入导出配置
# ═══════════════════════════════════════════════════════════════════════════════

_ROW_LIMIT = 500

# N4 各sheet导出配置
_SHEET_CONFIGS: dict[str, dict[str, Any]] = {
    "N4-1": {
        "title": "N4-1 审定表",
        "headers": [
            "税种", "本期发生额", "未审数", "审计调整(AJE)",
            "重分类调整(RJE)", "审定数", "上期发生额", "同比变动(%)",
            "N2计提额", "差异", "核查结论", "备注",
        ],
        "fields": [
            "taxType", "currentAmount", "unadjustedAmount", "ajeAmount",
            "rjeAmount", "auditedAmount", "priorAmount", "yoyChange",
            "n2AccrualAmount", "difference", "checkConclusion", "remark",
        ],
    },
    "N4-2": {
        "title": "N4-2 明细表",
        "headers": [
            "序号", "税种", "计税依据", "适用税率",
            "本期发生额", "上期发生额", "同比变动(%)",
            "N2计提额", "差异", "核查结论", "备注",
        ],
        "fields": [
            "seqNo", "taxType", "taxBase", "taxRate",
            "currentAmount", "priorAmount", "yoyChange",
            "n2AccrualAmount", "difference", "checkConclusion", "remark",
        ],
    },
    "N4-3": {
        "title": "N4-3 调整分录",
        "headers": [
            "调整类型", "科目编码", "科目名称", "摘要",
            "借方金额", "贷方金额", "备注",
        ],
        "fields": [
            "adjustmentType", "accountCode", "accountName", "summary",
            "debitAmount", "creditAmount", "remark",
        ],
    },
}

_NUMERIC_FIELDS: set[str] = {
    "currentAmount", "unadjustedAmount", "ajeAmount", "rjeAmount",
    "auditedAmount", "priorAmount", "yoyChange",
    "n2AccrualAmount", "difference",
    "taxBase", "taxRate",
    "debitAmount", "creditAmount",
}


# ═══════════════════════════════════════════════════════════════════════════════
# 导出辅助
# ═══════════════════════════════════════════════════════════════════════════════


def _style_header_row(ws: Any) -> None:
    """给表头行设置样式"""
    header_font = Font(bold=True, size=11)
    header_fill = PatternFill(start_color="D9E1F2", end_color="D9E1F2", fill_type="solid")
    header_align = Alignment(horizontal="center", vertical="center", wrap_text=True)

    for cell in ws[1]:
        cell.font = header_font
        cell.fill = header_fill
        cell.alignment = header_align
        ws.column_dimensions[cell.column_letter].width = max(
            len(str(cell.value or "")) * 2 + 4, 12
        )
    ws.freeze_panes = "A2"


def _parse_num(v: Any) -> float:
    """安全解析数值，None/空/NaN→0.0."""
    if v is None or v == "":
        return 0.0
    try:
        f = float(v)
        return 0.0 if f != f else f  # NaN→0.0
    except (ValueError, TypeError):
        return 0.0


def _safe_str(val: Any) -> str:
    if val is None:
        return ""
    return str(val).strip()


# ═══════════════════════════════════════════════════════════════════════════════
# 导出模板
# ═══════════════════════════════════════════════════════════════════════════════


async def build_template_workbook(
    wp_id: str, db: Any, sheet: str | None = None
) -> io.BytesIO:
    """生成空白xlsx模板（多sheet或单sheet）.

    Sheet列表: N4-1(审定表12列), N4-2(明细表11列), N4-3(调整分录7列)

    Requirements: 6.2
    """
    wb = Workbook()
    wb.remove(wb.active)

    configs_to_export = _SHEET_CONFIGS
    if sheet and sheet in _SHEET_CONFIGS:
        configs_to_export = {sheet: _SHEET_CONFIGS[sheet]}

    for _sheet_key, config in configs_to_export.items():
        ws = wb.create_sheet(title=config["title"][:31])
        for col_idx, header in enumerate(config["headers"], 1):
            ws.cell(row=1, column=col_idx, value=header)
        _style_header_row(ws)

    buffer = io.BytesIO()
    wb.save(buffer)
    buffer.seek(0)
    return buffer


# ═══════════════════════════════════════════════════════════════════════════════
# 导出数据
# ═══════════════════════════════════════════════════════════════════════════════


async def build_data_workbook(
    wp_id: str, db: Any, sheet: str | None = None
) -> io.BytesIO:
    """导出当前数据为xlsx（多sheet或单sheet）.

    从 checklist_responses 读取已保存的数据并填入xlsx。

    Requirements: 6.2
    """
    import sqlalchemy as sa

    wb = Workbook()
    wb.remove(wb.active)

    configs_to_export = _SHEET_CONFIGS
    if sheet and sheet in _SHEET_CONFIGS:
        configs_to_export = {sheet: _SHEET_CONFIGS[sheet]}

    for sheet_key, config in configs_to_export.items():
        ws = wb.create_sheet(title=config["title"][:31])
        for col_idx, header in enumerate(config["headers"], 1):
            ws.cell(row=1, column=col_idx, value=header)
        _style_header_row(ws)

        # 读取对应数据
        item_id = f"N4-{sheet_key}-rows"
        result = await db.execute(
            sa.text(
                "SELECT remark FROM checklist_responses "
                "WHERE wp_id = :wp_id AND item_id = :item_id LIMIT 1"
            ),
            {"wp_id": wp_id, "item_id": item_id},
        )
        row = result.fetchone()
        data_rows: list[dict] = []
        if row and row.remark:
            try:
                data_rows = json.loads(row.remark)
            except (json.JSONDecodeError, TypeError):
                pass

        for data_row in data_rows:
            row_values = [data_row.get(f, "") or "" for f in config["fields"]]
            ws.append(row_values)

    buffer = io.BytesIO()
    wb.save(buffer)
    buffer.seek(0)
    return buffer


# ═══════════════════════════════════════════════════════════════════════════════
# 导入数据
# ═══════════════════════════════════════════════════════════════════════════════


async def parse_import_file(
    wp_id: str, file_content: bytes, db: Any, sheet: str | None = None
) -> dict[str, Any]:
    """解析上传的xlsx并写入checklist_responses.

    支持多sheet导入（自动识别区段）或单sheet导入。
    验证表头结构匹配。

    Requirements: 6.2
    """
    import sqlalchemy as sa

    if len(file_content) > 10 * 1024 * 1024:
        raise ValueError("文件大小超过 10MB 限制")

    try:
        wb = load_workbook(io.BytesIO(file_content), read_only=True, data_only=True)
    except Exception:
        raise ValueError("无法解析xlsx文件")

    imported_counts: dict[str, int] = {}
    warnings: list[str] = []

    for ws in wb.worksheets:
        ws_title = ws.title.strip()

        # 尝试匹配到配置
        matched_config: dict[str, Any] | None = None
        matched_key: str | None = None

        for sheet_key, config in _SHEET_CONFIGS.items():
            if config["title"][:31] == ws_title or sheet_key in ws_title:
                matched_config = config
                matched_key = sheet_key
                break

        if not matched_config or not matched_key:
            continue

        # 若指定了sheet，跳过不匹配的
        if sheet and matched_key != sheet:
            continue

        # 读取表头
        actual_headers: list[str] = []
        for cell in next(ws.iter_rows(min_row=1, max_row=1)):
            if cell.value is not None:
                actual_headers.append(str(cell.value).strip())

        # 校验表头结构
        expected_headers = matched_config["headers"]
        if actual_headers and actual_headers[:3] != expected_headers[:3]:
            warnings.append(f"Sheet '{ws_title}' 表头结构不匹配，尝试按位置解析")

        # 解析数据行
        parsed_rows: list[dict] = []
        for row_idx, row in enumerate(ws.iter_rows(min_row=2, values_only=True), start=2):
            if row_idx > _ROW_LIMIT + 1:
                warnings.append(f"Sheet '{ws_title}' 超过 {_ROW_LIMIT} 行限制，已截断")
                break
            if all(cell is None or str(cell).strip() == "" for cell in row):
                continue

            row_data: dict[str, Any] = {"rowId": str(uuid4())}
            fields = matched_config["fields"]
            headers = matched_config["headers"]

            for i, (header, field) in enumerate(zip(headers, fields)):
                try:
                    header_idx = actual_headers.index(header)
                    raw_val = row[header_idx] if header_idx < len(row) else None
                except (ValueError, IndexError):
                    raw_val = row[i] if i < len(row) else None

                if field in _NUMERIC_FIELDS:
                    row_data[field] = _parse_num(raw_val)
                else:
                    row_data[field] = _safe_str(raw_val)

            parsed_rows.append(row_data)

        imported_counts[matched_key] = len(parsed_rows)

        # 写入 checklist_responses
        item_id = f"N4-{matched_key}-rows"
        json_str = json.dumps(parsed_rows, ensure_ascii=False)
        await db.execute(
            sa.text(
                "INSERT INTO checklist_responses (id, wp_id, item_id, remark) "
                "VALUES (:id, :wp_id, :item_id, :remark) "
                "ON CONFLICT (wp_id, item_id) DO UPDATE SET remark = :remark"
            ),
            {"id": str(uuid4()), "wp_id": wp_id, "item_id": item_id, "remark": json_str},
        )

    await db.flush()

    return {
        "imported_count": sum(imported_counts.values()),
        "imported_counts": imported_counts,
        "warnings": warnings,
    }

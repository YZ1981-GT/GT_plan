"""L8 财务费用 — 业务逻辑服务

纯函数计算引擎（无DB依赖，可单独测试）+ xlsx导入导出逻辑 + 损益类发生额查询。

提供功能：
- 损益类发生额（从tb_ledger取借方发生-贷方发生，非余额！）
- 利息汇总（L1短期借款/L3长期借款/L4应付债券/L5未确认融资费用摊销）
- 非金融机构利息测算（可扣除利息/超标利息）
- 截止测试提取（报告日±N天序时账）
- 导出模板 / 导出数据 / 导入数据

科目6603财务费用（**损益类/借方**）：本期发生额=借方发生-贷方发生
L筹资循环利息汇聚终点。

Requirements: 2.4, 4.5, 5.2-5.3, 6.2-6.4, 8.8
"""

from __future__ import annotations

import io
import json
import logging
from datetime import date, timedelta
from typing import Any
from uuid import uuid4

from openpyxl import Workbook, load_workbook
from openpyxl.styles import Alignment, Font, PatternFill

logger = logging.getLogger(__name__)


# ═══════════════════════════════════════════════════════════════════════════════
# 损益类发生额（核心纯函数）
# ═══════════════════════════════════════════════════════════════════════════════


def calc_occurrence(debit_occur: float, credit_occur: float) -> float:
    """损益类本期发生额 = 借方发生 - 贷方发生（费用类科目为借方）.

    Requirements: 2.4, 8.2
    """
    return debit_occur - credit_occur


def calc_net_finance_expense(
    interest_exp: float,
    interest_inc: float,
    fx: float,
    fee: float,
    other: float,
) -> float:
    """净财务费用 = 利息支出 - 利息收入 + 汇兑损益 + 手续费 + 其他.

    Requirements: 3.2, 8.3
    """
    return interest_exp - interest_inc + fx + fee + other


# ═══════════════════════════════════════════════════════════════════════════════
# 利息汇总（L1/L3/L4/L5联动）
# ═══════════════════════════════════════════════════════════════════════════════


def aggregate_interest(l1: float, l3: float, l4: float, l5: float) -> float:
    """汇总各来源利息测算合计（L1短期借款+L3长期借款+L4应付债券+L5摊销）.

    Requirements: 4.5, 8.5
    """
    return l1 + l3 + l4 + l5


def calc_interest_diff(estimated: float, booked: float) -> float:
    """测算利息vs账面差异.

    Requirements: 4.6
    """
    return estimated - booked


# ═══════════════════════════════════════════════════════════════════════════════
# 非金融机构利息测算（纯函数）
# ═══════════════════════════════════════════════════════════════════════════════


def calc_deductible_interest(principal: float, benchmark_rate: float, days: int) -> float:
    """可扣除利息 = 本金 × 同期金融机构利率 × 天数 / 360.

    Requirements: 5.2, 8.5
    """
    return principal * benchmark_rate * days / 360


def calc_excess_interest(booked: float, deductible: float) -> float:
    """超标利息 = 账载利息 - 可扣除利息.

    Requirements: 5.3, 8.6
    """
    return booked - deductible


def calc_non_fin_interest(rows: list[dict[str, Any]]) -> list[dict[str, Any]]:
    """批量计算非金融机构利息测算.

    每行 dict 需含: principal, agreed_rate, benchmark_rate, days, booked_interest
    返回增强后的行列表，附加 deductible_interest, excess_interest 字段。

    Requirements: 5.2-5.3
    """
    results: list[dict[str, Any]] = []
    for row in rows:
        principal = float(row.get("principal") or 0)
        benchmark_rate = float(row.get("benchmark_rate") or 0)
        days = int(row.get("days") or 0)
        booked = float(row.get("booked_interest") or 0)

        deductible = calc_deductible_interest(principal, benchmark_rate, days)
        excess = calc_excess_interest(booked, deductible)

        enhanced = {**row}
        enhanced["deductible_interest"] = round(deductible, 2)
        enhanced["excess_interest"] = round(excess, 2)
        enhanced["is_excess"] = excess > 0.01
        results.append(enhanced)
    return results


# ═══════════════════════════════════════════════════════════════════════════════
# 损益类发生额查询（DB依赖）
# ═══════════════════════════════════════════════════════════════════════════════


async def get_occurrence_amount(
    db: Any,
    project_id: Any,
    year: int,
    account_code: str = "6603",
) -> dict[str, Any]:
    """从 tb_ledger 取损益类科目的借方/贷方发生额汇总.

    损益类科目取发生额（非余额！），从序时账聚合。
    使用 account_code LIKE '{code}%' 包含所有明细科目。

    Requirements: 2.4
    """
    import sqlalchemy as sa

    debit_total = 0.0
    credit_total = 0.0
    try:
        result = await db.execute(
            sa.text("""
                SELECT
                    COALESCE(SUM(debit_amount), 0) AS debit_total,
                    COALESCE(SUM(credit_amount), 0) AS credit_total
                FROM tb_ledger
                WHERE project_id = :pid
                  AND year = :year
                  AND account_code LIKE :code_prefix
                  AND is_deleted = false
            """),
            {
                "pid": str(project_id),
                "year": year,
                "code_prefix": f"{account_code}%",
            },
        )
        row = result.fetchone()
        if row:
            debit_total = float(row.debit_total or 0)
            credit_total = float(row.credit_total or 0)
    except Exception as e:  # noqa: BLE001
        logger.warning("L8 service: tb_ledger 发生额查询失败 account=%s: %s", account_code, e)

    occurrence = debit_total - credit_total
    return {
        "account_code": account_code,
        "debit_total": round(debit_total, 2),
        "credit_total": round(credit_total, 2),
        "occurrence": round(occurrence, 2),
    }


# ═══════════════════════════════════════════════════════════════════════════════
# 利息汇总查询（从L1/L3/L4/L5底稿checklist_responses取）
# ═══════════════════════════════════════════════════════════════════════════════


async def aggregate_interest_from_l_cycle(
    db: Any,
    project_id: Any,
) -> dict[str, Any]:
    """从 L1/L3/L4/L5 底稿的 checklist_responses 读取利息测算值.

    查找各底稿已计算的利息测算值（存储在 conclusion/remark JSON 中）。
    汇总各来源利息 → 测算利息支出合计。

    Requirements: 4.5
    """
    import sqlalchemy as sa

    interest_sources: dict[str, float] = {
        "l1_short_term": 0.0,
        "l3_long_term": 0.0,
        "l4_bonds": 0.0,
        "l5_amortization": 0.0,
    }

    # 查找L1/L3/L4/L5的利息测算汇总
    # 各底稿在保存时会把利息测算值存储在特定 item_id
    source_patterns = {
        "l1_short_term": "L1-interest-total%",
        "l3_long_term": "L3-interest-total%",
        "l4_bonds": "L4-interest-total%",
        "l5_amortization": "L5-amortization-total%",
    }

    try:
        # 通过 wp_index 找到同项目同循环的 L1/L3/L4/L5 底稿
        wp_result = await db.execute(
            sa.text("""
                SELECT wp.id, wi.wp_code
                FROM working_papers wp
                JOIN wp_index wi ON wi.id = wp.wp_index_id
                WHERE wp.project_id = :pid
                  AND wi.wp_code = ANY(:codes)
                  AND wp.is_deleted = false
            """),
            {
                "pid": str(project_id),
                "codes": ["L1-2", "L3-2", "L4-2", "L5-5"],
            },
        )
        wp_rows = wp_result.fetchall()

        for wp_row in wp_rows:
            wp_id = str(wp_row.id)
            wp_code = wp_row.wp_code

            # 从各底稿的 checklist_responses 取利息合计值
            cr_result = await db.execute(
                sa.text("""
                    SELECT item_id, remark
                    FROM checklist_responses
                    WHERE wp_id = :wp_id
                      AND (item_id LIKE :pattern_interest OR item_id LIKE :pattern_total)
                    LIMIT 10
                """),
                {
                    "wp_id": wp_id,
                    "pattern_interest": f"{wp_code[:2]}%-interest-total%",
                    "pattern_total": f"{wp_code[:2]}%-total%",
                },
            )

            for cr_row in cr_result.fetchall():
                amount = 0.0
                if cr_row.remark:
                    try:
                        parsed = json.loads(cr_row.remark)
                        if isinstance(parsed, dict):
                            amount = float(
                                parsed.get("interest_total")
                                or parsed.get("total")
                                or parsed.get("amount")
                                or 0
                            )
                    except (json.JSONDecodeError, TypeError, ValueError):
                        pass

                if wp_code.startswith("L1"):
                    interest_sources["l1_short_term"] = max(interest_sources["l1_short_term"], amount)
                elif wp_code.startswith("L3"):
                    interest_sources["l3_long_term"] = max(interest_sources["l3_long_term"], amount)
                elif wp_code.startswith("L4"):
                    interest_sources["l4_bonds"] = max(interest_sources["l4_bonds"], amount)
                elif wp_code.startswith("L5"):
                    interest_sources["l5_amortization"] = max(interest_sources["l5_amortization"], amount)

    except Exception as e:  # noqa: BLE001
        logger.warning("L8 service: L循环利息汇总查询失败: %s", e)

    total = aggregate_interest(
        interest_sources["l1_short_term"],
        interest_sources["l3_long_term"],
        interest_sources["l4_bonds"],
        interest_sources["l5_amortization"],
    )

    return {
        "sources": interest_sources,
        "total_estimated": round(total, 2),
    }


# ═══════════════════════════════════════════════════════════════════════════════
# 截止测试提取（报告日±N天序时账）
# ═══════════════════════════════════════════════════════════════════════════════


async def extract_cutoff_entries(
    db: Any,
    project_id: Any,
    report_date: date,
    days: int = 5,
    account_code: str = "6603",
) -> list[dict[str, Any]]:
    """从 tb_ledger 提取报告日前后±N天的财务费用序时账明细.

    截止测试核心逻辑：期末前后费用归属期间正确性验证。

    Requirements: 6.2-6.4
    """
    import sqlalchemy as sa

    date_start = report_date - timedelta(days=days)
    date_end = report_date + timedelta(days=days)

    entries: list[dict[str, Any]] = []
    try:
        result = await db.execute(
            sa.text("""
                SELECT
                    voucher_no, voucher_date, summary,
                    debit_amount, credit_amount, account_code,
                    accounting_period
                FROM tb_ledger
                WHERE project_id = :pid
                  AND account_code LIKE :code_prefix
                  AND voucher_date BETWEEN :date_start AND :date_end
                  AND is_deleted = false
                ORDER BY voucher_date, voucher_no
                LIMIT 500
            """),
            {
                "pid": str(project_id),
                "code_prefix": f"{account_code}%",
                "date_start": date_start.isoformat(),
                "date_end": date_end.isoformat(),
            },
        )

        for row in result.fetchall():
            voucher_date = row.voucher_date
            # 判断是否跨期：入账期间 vs 归属期间
            booking_period = str(row.accounting_period or "")
            # 归属期间基于报告日判定
            attribution_period = (
                str(report_date.year)
                if voucher_date and voucher_date <= report_date
                else str(report_date.year + 1) if voucher_date else ""
            )
            is_cross_period = booking_period != "" and attribution_period != booking_period

            entries.append({
                "voucher_no": row.voucher_no or "",
                "voucher_date": str(voucher_date) if voucher_date else "",
                "summary": row.summary or "",
                "debit_amount": float(row.debit_amount or 0),
                "credit_amount": float(row.credit_amount or 0),
                "account_code": row.account_code or "",
                "booking_period": booking_period,
                "attribution_period": attribution_period,
                "is_cross_period": is_cross_period,
            })
    except Exception as e:  # noqa: BLE001
        logger.warning("L8 service: 截止测试提取失败: %s", e)

    return entries


# ═══════════════════════════════════════════════════════════════════════════════
# 导入导出配置
# ═══════════════════════════════════════════════════════════════════════════════

_ROW_LIMIT = 500

# L8 各sheet导出配置
_SHEET_CONFIGS: dict[str, dict[str, Any]] = {
    "L8-1": {
        "title": "L8-1 审定表",
        "headers": [
            "项目", "本期发生额", "上期发生额", "变动",
            "未审数", "审计调整", "重分类调整", "审定数", "备注",
        ],
        "fields": [
            "itemName", "currentAmount", "priorAmount", "changeAmount",
            "unadjustedAmount", "ajeAmount", "rjeAmount", "auditedAmount", "remark",
        ],
    },
    "L8-2": {
        "title": "L8-2 明细表",
        "headers": [
            "费用项目", "本期发生额", "上期发生额", "变动",
            "变动率(%)", "说明",
        ],
        "fields": [
            "itemName", "currentAmount", "priorAmount", "changeAmount",
            "changeRate", "remark",
        ],
    },
    "L8-3": {
        "title": "L8-3 调整分录",
        "headers": [
            "调整类型", "科目编码", "科目名称", "摘要",
            "借方金额", "贷方金额", "备注",
        ],
        "fields": [
            "adjustmentType", "accountCode", "accountName", "summary",
            "debitAmount", "creditAmount", "remark",
        ],
    },
    "L8-4": {
        "title": "L8-4 非金融机构利息测算",
        "headers": [
            "借款方", "本金", "约定利率(%)", "同期金融机构利率(%)",
            "天数", "账载利息", "可扣除利息", "超标利息", "备注",
        ],
        "fields": [
            "borrower", "principal", "agreedRate", "benchmarkRate",
            "days", "bookedInterest", "deductibleInterest", "excessInterest", "remark",
        ],
    },
    "L8-5": {
        "title": "L8-5 截止测试",
        "headers": [
            "凭证号", "日期", "摘要", "借方金额", "贷方金额",
            "入账期间", "应归属期间", "是否跨期", "检查结论", "备注",
        ],
        "fields": [
            "voucherNo", "voucherDate", "summary", "debitAmount", "creditAmount",
            "bookingPeriod", "attributionPeriod", "isCrossPeriod", "checkResult", "remark",
        ],
    },
}

_NUMERIC_FIELDS: set[str] = {
    "currentAmount", "priorAmount", "changeAmount", "changeRate",
    "unadjustedAmount", "ajeAmount", "rjeAmount", "auditedAmount",
    "debitAmount", "creditAmount", "principal", "agreedRate",
    "benchmarkRate", "days", "bookedInterest", "deductibleInterest", "excessInterest",
}


# ═══════════════════════════════════════════════════════════════════════════════
# 导出辅助
# ═══════════════════════════════════════════════════════════════════════════════


def _style_header_row(ws: Any) -> None:
    """给表头行设置样式"""
    header_font = Font(bold=True, size=11)
    header_fill = PatternFill(start_color="D9E1F2", end_color="D9E1F2", fill_type="solid")
    header_align = Alignment(horizontal="center", vertical="center", wrap_text=True)

    for cell in ws[1]:
        cell.font = header_font
        cell.fill = header_fill
        cell.alignment = header_align
        ws.column_dimensions[cell.column_letter].width = max(
            len(str(cell.value or "")) * 2 + 4, 12
        )
    ws.freeze_panes = "A2"


def _safe_float(val: Any) -> float:
    if val is None:
        return 0.0
    try:
        return float(val)
    except (ValueError, TypeError):
        return 0.0


def _safe_str(val: Any) -> str:
    if val is None:
        return ""
    return str(val).strip()


# ═══════════════════════════════════════════════════════════════════════════════
# 导出模板
# ═══════════════════════════════════════════════════════════════════════════════


async def export_template(wp_id: str, db: Any, sheet: str | None = None) -> io.BytesIO:
    """生成空白xlsx模板（多sheet或单sheet）.

    Requirements: 8.8
    """
    wb = Workbook()
    wb.remove(wb.active)

    configs_to_export = _SHEET_CONFIGS
    if sheet and sheet in _SHEET_CONFIGS:
        configs_to_export = {sheet: _SHEET_CONFIGS[sheet]}

    for _sheet_key, config in configs_to_export.items():
        ws = wb.create_sheet(title=config["title"][:31])
        for col_idx, header in enumerate(config["headers"], 1):
            ws.cell(row=1, column=col_idx, value=header)
        _style_header_row(ws)

    buffer = io.BytesIO()
    wb.save(buffer)
    buffer.seek(0)
    return buffer


# ═══════════════════════════════════════════════════════════════════════════════
# 导出数据
# ═══════════════════════════════════════════════════════════════════════════════


async def export_data(wp_id: str, db: Any, sheet: str | None = None) -> io.BytesIO:
    """导出当前数据为xlsx（多sheet或单sheet）.

    从 checklist_responses 读取已保存的数据并填入xlsx。

    Requirements: 8.8
    """
    import sqlalchemy as sa

    wb = Workbook()
    wb.remove(wb.active)

    configs_to_export = _SHEET_CONFIGS
    if sheet and sheet in _SHEET_CONFIGS:
        configs_to_export = {sheet: _SHEET_CONFIGS[sheet]}

    for sheet_key, config in configs_to_export.items():
        ws = wb.create_sheet(title=config["title"][:31])
        for col_idx, header in enumerate(config["headers"], 1):
            ws.cell(row=1, column=col_idx, value=header)
        _style_header_row(ws)

        # 读取对应数据
        item_id = f"L8-{sheet_key}-rows"
        result = await db.execute(
            sa.text(
                "SELECT remark FROM checklist_responses "
                "WHERE wp_id = :wp_id AND item_id = :item_id LIMIT 1"
            ),
            {"wp_id": wp_id, "item_id": item_id},
        )
        row = result.fetchone()
        data_rows: list[dict] = []
        if row and row.remark:
            try:
                data_rows = json.loads(row.remark)
            except (json.JSONDecodeError, TypeError):
                pass

        for data_row in data_rows:
            row_values = [data_row.get(f, "") or "" for f in config["fields"]]
            ws.append(row_values)

    buffer = io.BytesIO()
    wb.save(buffer)
    buffer.seek(0)
    return buffer


# ═══════════════════════════════════════════════════════════════════════════════
# 导入数据
# ═══════════════════════════════════════════════════════════════════════════════


async def import_data(
    wp_id: str, file_content: bytes, db: Any, sheet: str | None = None
) -> dict[str, Any]:
    """解析上传的xlsx并写入checklist_responses.

    支持多sheet导入（自动识别区段）或单sheet导入。

    Requirements: 8.8
    """
    import sqlalchemy as sa

    if len(file_content) > 10 * 1024 * 1024:
        raise ValueError("文件大小超过 10MB 限制")

    try:
        wb = load_workbook(io.BytesIO(file_content), read_only=True, data_only=True)
    except Exception:
        raise ValueError("无法解析xlsx文件")

    imported_counts: dict[str, int] = {}
    warnings: list[str] = []

    for ws in wb.worksheets:
        ws_title = ws.title.strip()

        # 尝试匹配到配置
        matched_config: dict[str, Any] | None = None
        matched_key: str | None = None

        for sheet_key, config in _SHEET_CONFIGS.items():
            if config["title"][:31] == ws_title or sheet_key in ws_title:
                matched_config = config
                matched_key = sheet_key
                break

        if not matched_config or not matched_key:
            continue

        # 若指定了sheet，跳过不匹配的
        if sheet and matched_key != sheet:
            continue

        # 读取表头
        actual_headers: list[str] = []
        for cell in next(ws.iter_rows(min_row=1, max_row=1)):
            if cell.value is not None:
                actual_headers.append(str(cell.value).strip())

        # 校验表头结构
        expected_headers = matched_config["headers"]
        if actual_headers and actual_headers[:3] != expected_headers[:3]:
            warnings.append(f"Sheet '{ws_title}' 表头结构不匹配，尝试按位置解析")

        # 解析数据行
        parsed_rows: list[dict] = []
        for row_idx, row in enumerate(ws.iter_rows(min_row=2, values_only=True), start=2):
            if row_idx > _ROW_LIMIT + 1:
                warnings.append(f"Sheet '{ws_title}' 超过 {_ROW_LIMIT} 行限制，已截断")
                break
            if all(cell is None or str(cell).strip() == "" for cell in row):
                continue

            row_data: dict[str, Any] = {"rowId": str(uuid4())}
            fields = matched_config["fields"]
            headers = matched_config["headers"]

            for i, (header, field) in enumerate(zip(headers, fields)):
                try:
                    header_idx = actual_headers.index(header)
                    raw_val = row[header_idx] if header_idx < len(row) else None
                except (ValueError, IndexError):
                    raw_val = row[i] if i < len(row) else None

                if field in _NUMERIC_FIELDS:
                    row_data[field] = _safe_float(raw_val)
                else:
                    row_data[field] = _safe_str(raw_val)

            parsed_rows.append(row_data)

        imported_counts[matched_key] = len(parsed_rows)

        # 写入 checklist_responses
        item_id = f"L8-{matched_key}-rows"
        json_str = json.dumps(parsed_rows, ensure_ascii=False)
        await db.execute(
            sa.text(
                "INSERT INTO checklist_responses (id, wp_id, item_id, remark) "
                "VALUES (:id, :wp_id, :item_id, :remark) "
                "ON CONFLICT (wp_id, item_id) DO UPDATE SET remark = :remark"
            ),
            {"id": str(uuid4()), "wp_id": wp_id, "item_id": item_id, "remark": json_str},
        )

    await db.flush()

    return {
        "imported_count": sum(imported_counts.values()),
        "imported_counts": imported_counts,
        "warnings": warnings,
    }

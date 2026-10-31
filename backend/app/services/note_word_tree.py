"""四版 Word 附注模板 ↔ 种子 JSON：章节树与表格比对（只读，不连库）。

spec: disclosure-note-template-word-realignment Task 1（design §2）

权威 = ``data/audit_report_templates/disclosure_notes/{variant}.docx``。口径：

- 章 = Heading 1；节 = 国企 Heading 2∪3、上市 Heading 2（``note_section_kind.SECTION_HEADING_LEVELS``），
  更深的标题是节内小标题；
- 表格归属：节标题之后、下一个节或章标题之前；章标题与首个节标题之间的表格属章正文；
- 标题比对用 ``note_section_kind.normalize_title``；表头 / 行标签去空白、统一全半角括号与冒号后比对；
- 合并单元格按 OOXML ``gridSpan`` / ``vMerge`` 解析，表头行数判定与
  ``scripts/seed/sync_note_templates_from_word.detect_header_row_count`` 同一规则。

请求期不调用本模块（docx 延迟导入）。调用方：``scripts/check/check_note_template_word_tree.py``、
``scripts/seed/realign_note_templates_to_word.py`` 与守卫 ``test_note_template_word_tree_alignment``。
"""

from __future__ import annotations

import difflib
import json
import re
from collections import defaultdict
from dataclasses import dataclass, field
from pathlib import Path
from typing import Any

from app.services.note_section_catalog import (
    DATA_DIR,
    VARIANT_KEYS,
    WORD_TEMPLATE_DIR,
    section_applies_to_scope,
    template_json_path,
)
from app.services.note_section_kind import (
    SECTION_HEADING_LEVELS,
    WORD_TITLE_ALIASES,
    normalize_title,
)
from app.services.note_title_hint import split_title_hint
from app.services.note_word_table import (
    WordTable,
    compare_tables,
    iter_blocks,
    norm_cell,
    parse_word_table,
)

__all__ = [
    "Diff",
    "SeedChapter",
    "SeedSection",
    "VariantMatch",
    "WordChapter",
    "WordDoc",
    "WordSection",
    "WordTable",
    "ALLOWLIST_PATH",
    "build_report",
    "cached_word_variant",
    "check_against_allowlist",
    "clean_word_heading",
    "compare_tables",
    "compare_variant",
    "load_seed",
    "load_word_variant",
    "match_all",
    "match_variant",
    "norm_cell",
    "parse_word_docx",
    "seed_diffs",
    "split_title_hint",
]


#: 白名单：``[{variant, kind, key, reason}]``，理由必填；只许放「与 Word 不同是有意为之」的条目
ALLOWLIST_PATH = DATA_DIR / "note_template_word_alignment_allowlist.json"




# ── 数据结构 ───────────────────────────────────────────────────────────────




@dataclass
class WordSection:
    title: str
    norm: str
    level: int
    body: list[str] = field(default_factory=list)  # 节标题后、首张表前的非标题段落
    subtitles: list[str] = field(default_factory=list)
    tables: list[WordTable] = field(default_factory=list)
    blocks: list[tuple[int | None, str]] = field(default_factory=list)  # 全部文字块（小标题层级 / None=段落），按文档顺序


@dataclass
class WordChapter:
    index: int  # H1 序号（1 起）
    title: str
    norm: str
    body: list[str] = field(default_factory=list)
    tables: list[WordTable] = field(default_factory=list)  # 章正文表
    sections: list[WordSection] = field(default_factory=list)


@dataclass
class WordDoc:
    variant: str
    chapters: list[WordChapter]


@dataclass
class SeedSection:
    code: str
    title: str
    norm: str
    scope: str | None
    section_id: str | None
    chapter_code: str | None
    position: int  # 全种子排序后的位置
    raw: dict[str, Any]
    parent_ok: bool = True  # parent_section_id 能解析到章；False 时按编号前缀归章
    key: str = ""  # 报告键 = 编号；同号第 n 个（n≥2）为「编号#n」


@dataclass
class SeedChapter:
    code: str
    title: str
    norm: str
    scope: str | None
    section_id: str | None
    position: int
    raw: dict[str, Any]
    sections: list[SeedSection] = field(default_factory=list)

    def scope_title(self, report_scope: str) -> str:
        """口径章名：``scope_titles[report_scope]``，无则用 ``section_title``（上市项目附注章单体 / 合并不同名）。"""
        return (self.raw.get("scope_titles") or {}).get(report_scope) or self.title

    def scope_norm(self, report_scope: str) -> str:
        return normalize_title(self.scope_title(report_scope))


@dataclass(frozen=True, order=True)
class Diff:
    variant: str
    kind: str
    key: str
    detail: str = ""

    def as_dict(self) -> dict[str, str]:
        return {"variant": self.variant, "kind": self.kind, "key": self.key, "detail": self.detail}


# ── 文字规范化 ─────────────────────────────────────────────────────────────





# ── Word 表格（OOXML 合并单元格）────────────────────────────────────────────












# ── Word 文档 → 章节树 ─────────────────────────────────────────────────────


#: 模板标记（``##SECTION:…##`` / ``{{seq:七}}`` / ``{{section:…}}`` / ``{{table:…}}``）不是内容
_MARKER = re.compile(r"##[^#]*##|\{\{[^{}]*\}\}")
#: Word 模板的占位序号「X、数据资源」（真实附注里会换成数字）
_PLACEHOLDER_ENUM = re.compile(r"^[XxＸｘ][、.．]")


def clean_word_heading(text: str) -> str:
    """Word 标题原文去模板占位序号「X、」（种子标题用它）。"""
    return _PLACEHOLDER_ENUM.sub("", text.strip()).strip()


def word_title_norm(text: str) -> str:
    return normalize_title(clean_word_heading(text))


def _chapter_events(path: Path | str) -> list[tuple[str, list[tuple[str, int | None, Any]]]]:
    """按 H1 切章：[(章标题, [("h", 层级, 文字) | ("p", None, 文字) | ("t", None, WordTable)])]。

    标题层级与上传解析同一读法（``note_document_parser.docx_paragraph_level``：样式名 → 段落大纲级别 → 样式链），
    再按编号排位嵌套同级标题（``nest_by_numbering``：「X、数据资源」后的「（1）」「（2）」「（3）」是它的组成部分）。
    """
    import docx
    from docx.text.paragraph import Paragraph

    from app.services.note_document_parser import docx_paragraph_heading, nest_by_numbering

    document = docx.Document(str(path))
    flat: list[tuple[str, int | None, Any]] = []
    outline_only: dict[int, bool] = {}
    style_levels: dict[str, int | None] = {}
    last_text = ""
    for tag, el in iter_blocks(document.element.body):
        if tag == "tbl":
            flat.append(("t", None, parse_word_table(el, last_text)))
            last_text = ""
            continue
        paragraph = Paragraph(el, document)
        text = _MARKER.sub("", paragraph.text or "").strip()
        if not text:
            continue
        level, from_style = docx_paragraph_heading(paragraph, el, style_levels)
        if level is not None:
            outline_only[len(flat)] = not from_style
        flat.append(("h", level, text) if level is not None else ("p", None, text))
        last_text = text
    heads = [k for k, (kind, _lvl, _x) in enumerate(flat) if kind == "h"]
    nested = nest_by_numbering([flat[k][1] for k in heads], [flat[k][2] for k in heads],
                               [outline_only.get(k, False) for k in heads])
    for k, level in zip(heads, nested):
        flat[k] = ("h", level, flat[k][2])
    chapters: list[tuple[str, list[tuple[str, int | None, Any]]]] = []
    for kind, level, payload in flat:
        if kind == "h" and level == 1:
            chapters.append((payload, []))
        elif chapters:
            chapters[-1][1].append((kind, level, payload))
    return chapters


def parse_word_docx(path: Path | str, std: str, *, variant: str = "") -> WordDoc:
    """解析一份附注 Word：章 / 节 / 小标题 / 正文段落 / 表格归属。

    国企节层级逐章判定：取本章出现的最高一级（H2 或 H3）。第七章「企业合并及合并财务报表」节是 H2，
    其下 H3（少数股东 / 主要财务信息 …）是节内小标题；第四、八章节是 H3。上市节恒为 H2。
    """
    allowed = SECTION_HEADING_LEVELS[std]
    out: list[WordChapter] = []
    for title, events in _chapter_events(path):
        chapter = WordChapter(index=len(out) + 1, title=title, norm=word_title_norm(title))
        present = sorted({lvl for kind, lvl, _ in events if kind == "h" and lvl in allowed})
        level = present[0] if present else None
        section: WordSection | None = None
        for kind, lvl, payload in events:
            holder = section if section is not None else chapter
            if kind == "t":
                holder.tables.append(payload)
            elif kind == "h" and lvl == level:
                section = WordSection(title=payload, norm=word_title_norm(payload), level=lvl)
                chapter.sections.append(section)
            elif kind == "h" and section is not None:
                section.subtitles.append(payload)
                section.blocks.append((lvl, payload))
            else:
                if isinstance(holder, WordSection):
                    holder.blocks.append((None, payload))
                if not holder.tables:  # 正文只收首张表之前的段落
                    holder.body.append(payload)
        out.append(chapter)
    return WordDoc(variant=variant, chapters=out)


def load_word_variant(variant: str) -> WordDoc:
    if variant not in VARIANT_KEYS:
        raise ValueError(f"未知口径：{variant}")
    return parse_word_docx(WORD_TEMPLATE_DIR / f"{variant}.docx", variant.split("_", 1)[0], variant=variant)


# ── 种子 JSON → 章节树 ─────────────────────────────────────────────────────


def _sort_key(item: tuple[int, dict[str, Any]]) -> tuple[float, int]:
    idx, sec = item
    order = sec.get("sort_order")
    return (float(order) if isinstance(order, (int, float)) else float("inf"), idx)


def load_seed(data: dict[str, Any] | str) -> list[SeedChapter]:
    """种子章节树（不按 scope 过滤）。``data`` 可为种子 dict 或准则名 soe / listed。"""
    if isinstance(data, str):
        data = json.loads(template_json_path(data).read_text(encoding="utf-8"))
    ordered = [s for _i, s in sorted(enumerate(data.get("sections") or []), key=_sort_key)]
    chapters: list[SeedChapter] = []
    by_id: dict[str, SeedChapter] = {}
    for pos, sec in enumerate(ordered):
        code = (sec.get("section_number") or "").strip()
        title = sec.get("section_title") or ""
        if sec.get("level") == 1:
            chapter = SeedChapter(
                code=code, title=title, norm=normalize_title(title), scope=sec.get("scope"),
                section_id=sec.get("section_id"), position=pos, raw=sec,
            )
            chapters.append(chapter)
            if chapter.section_id:
                by_id[chapter.section_id] = chapter
    seen_codes: dict[str, int] = defaultdict(int)
    for pos, sec in enumerate(ordered):
        if sec.get("level") == 1:
            continue
        parent = by_id.get(sec.get("parent_section_id") or "")
        code = (sec.get("section_number") or "").strip()
        title = sec.get("section_title") or ""
        seen_codes[code] += 1
        item = SeedSection(
            code=code, title=title, norm=normalize_title(title), scope=sec.get("scope"),
            section_id=sec.get("section_id"), chapter_code=parent.code if parent else None,
            position=pos, raw=sec, key=code if seen_codes[code] == 1 else f"{code}#{seen_codes[code]}",
        )
        if parent is not None:
            parent.sections.append(item)
        else:  # 父节点缺失：挂到编号前缀所属章，便于比对仍能定位
            prefix = code.split("、", 1)[0]
            host = next((c for c in chapters if c.code == prefix), None)
            item.parent_ok = False
            if host is not None:
                item.chapter_code = host.code
                host.sections.append(item)
    return chapters


# ── 配对与匹配 ─────────────────────────────────────────────────────────────


@dataclass
class VariantMatch:
    variant: str
    report_scope: str
    word: WordDoc
    chapters: list[SeedChapter]
    chapter_pairs: dict[int, SeedChapter]  # Word 章序 → 种子章
    matches: dict[int, tuple[WordChapter, WordSection]]  # 种子节 position → Word（章, 节）
    word_unmatched: list[tuple[WordChapter, WordSection]]
    body_leaves: dict[int, WordChapter]  # 章正文叶子「X、1 章名」：种子节 position → Word 章

    def sections_by_position(self) -> dict[int, SeedSection]:
        return {s.position: s for c in self.chapters for s in c.sections}


def _overlap(wc: WordChapter, sc: SeedChapter) -> int:
    counts: dict[str, int] = defaultdict(int)
    for s in sc.sections:
        counts[s.norm] += 1
    hit = 0
    for ws in wc.sections:
        if ws.norm and counts.get(ws.norm, 0) > 0:
            counts[ws.norm] -= 1
            hit += 1
    return hit


def _pair_chapters(word: WordDoc, chapters: list[SeedChapter], scope: str = "") -> dict[int, SeedChapter]:
    """章配对：标题相同（含口径章名）→ Word↔系统章名别名 → 章正文叶子用 Word 章名 → 对上的节最多。"""
    pairs: dict[int, SeedChapter] = {}
    used: set[int] = set()

    def take(wc: WordChapter, sc: SeedChapter | None) -> None:
        if sc is not None and wc.index not in pairs and sc.position not in used:
            pairs[wc.index] = sc
            used.add(sc.position)

    for wc in word.chapters:
        take(wc, next((c for c in chapters if c.position not in used and wc.norm in (c.norm, c.scope_norm(scope))), None))
    for wc in word.chapters:  # 章名别名：两侧都折叠后比（旧系统章名 ↔ Word 章名，方向无关）
        key = WORD_TITLE_ALIASES.get(wc.norm, wc.norm)
        take(wc, next((c for c in chapters if c.position not in used
                       and WORD_TITLE_ALIASES.get(c.norm, c.norm) == key), None))
    for wc in word.chapters:  # 种子章名错、但章正文叶子「X、1 章名」用的是 Word 章名（国企十三 / 十四、上市十）
        take(wc, next((
            c for c in chapters
            if c.position not in used and c.sections and c.sections[0].norm == wc.norm
        ), None))
    scored = sorted(
        (-_overlap(wc, sc), wc.index, sc.position, wc, sc)
        for wc in word.chapters if wc.index not in pairs
        for sc in chapters if sc.position not in used
    )
    for neg, _i, _p, wc, sc in scored:
        if neg < 0:
            take(wc, sc)
    return pairs


def _similar(a: str, b: str) -> bool:
    """同一空档内两个标题是否同一节：较短一方（≥4 字）被另一方包含，或字符相似度 ≥ 0.8。

    阈值实测：「营业收入与营业成本 / 营业收入和营业成本」0.89 命中；「其他应收款 / 应收账款」0.67 不命中。
    """
    if not a or not b:
        return False
    short, long_ = sorted((a, b), key=len)
    if len(short) >= 4 and short in long_:
        return True
    return difflib.SequenceMatcher(None, a, b).ratio() >= 0.8


def _gap_align(wc: WordChapter, sc: SeedChapter, matches: dict, seed_of: dict,
               leaf_candidates: set[int], bind: Any) -> None:
    """以本章已对上的节为锚点，两侧锚点之间的未匹配节按序找标题相似的一一对应（保持顺序，不交叉）。"""
    seed_seq = list(sc.sections)
    seed_idx = {s.position: i for i, s in enumerate(seed_seq)}
    anchors = [
        (wi, seed_idx[seed_of[id(ws)].position])
        for wi, ws in enumerate(wc.sections)
        if id(ws) in seed_of and seed_of[id(ws)].position in seed_idx
    ]
    drop = set(_out_of_order([(si, str(k)) for k, (_wi, si) in enumerate(anchors)]))
    anchors = [a for k, a in enumerate(anchors) if str(k) not in drop]
    bounds = [(-1, -1), *anchors, (len(wc.sections), len(seed_seq))]
    for (w0, s0), (w1, s1) in zip(bounds, bounds[1:]):
        w_gap = [ws for ws in wc.sections[w0 + 1:w1] if id(ws) not in seed_of and ws.norm]
        s_gap = [
            s for s in seed_seq[s0 + 1:s1]
            if s.position not in matches and s.position not in leaf_candidates
        ]
        cursor = 0
        for s in s_gap:
            hit = next((k for k in range(cursor, len(w_gap)) if _similar(w_gap[k].norm, s.norm)), None)
            if hit is not None:
                bind(s, wc, w_gap[hit])
                cursor = hit + 1


def match_variant(word: WordDoc, chapters: list[SeedChapter], variant: str) -> VariantMatch:
    """节匹配：先在已配对章内按标题依次对应，再跨章找同名（放错章）；剩余种子节判章正文叶子。"""
    report_scope = variant.split("_", 1)[1] if "_" in variant else "standalone"
    pairs = _pair_chapters(word, chapters, report_scope)
    matches: dict[int, tuple[WordChapter, WordSection]] = {}
    seed_of: dict[int, SeedSection] = {}  # id(WordSection) → 种子节

    def bind(s: SeedSection, wc: WordChapter, ws: WordSection) -> None:
        matches[s.position] = (wc, ws)
        seed_of[id(ws)] = s

    for wc in word.chapters:  # ① 已配对章内同名
        sc = pairs.get(wc.index)
        if sc is None:
            continue
        pool: dict[str, list[SeedSection]] = defaultdict(list)
        for s in sc.sections:
            pool[s.norm].append(s)
        for ws in wc.sections:
            if ws.norm and pool.get(ws.norm):
                bind(pool[ws.norm].pop(0), wc, ws)

    rest: dict[str, list[SeedSection]] = defaultdict(list)  # ② 跨章同名（放错章）
    for c in chapters:
        for s in c.sections:
            if s.position not in matches:
                rest[s.norm].append(s)
    for lst in rest.values():  # 同名多节：优先本口径适用的
        lst.sort(key=lambda s: (not section_applies_to_scope({"scope": s.scope}, report_scope), s.position))
    for wc in word.chapters:
        for ws in wc.sections:
            if id(ws) not in seed_of and ws.norm and rest.get(ws.norm):
                bind(rest[ws.norm].pop(0), wc, ws)

    by_code = {sc.code: word.chapters[idx - 1] for idx, sc in pairs.items()}
    leaf_candidates = {
        s.position
        for c in chapters if c.code in by_code
        for s in c.sections
        if s.position not in matches and s.norm and s.norm in (c.norm, by_code[c.code].norm)
    }
    for idx, sc in pairs.items():  # ③ 同章同位置的空档内一一对应（标题截断 / 措辞不同）
        _gap_align(word.chapters[idx - 1], sc, matches, seed_of, leaf_candidates, bind)

    body_leaves = {
        s.position: by_code[c.code]
        for c in chapters if c.code in by_code
        for s in c.sections if s.position in leaf_candidates and s.position not in matches
    }
    unmatched = [(wc, ws) for wc in word.chapters for ws in wc.sections if id(ws) not in seed_of]
    return VariantMatch(
        variant=variant, report_scope=report_scope, word=word, chapters=chapters,
        chapter_pairs=pairs, matches=matches, word_unmatched=unmatched, body_leaves=body_leaves,
    )


# ── 差异 ───────────────────────────────────────────────────────────────────


def _out_of_order(seq: list[tuple[int, str]]) -> list[str]:
    """seq = Word 顺序下的（种子位置, 键）；返回不在最长递增子序列中的键。"""
    if not seq:
        return []
    n = len(seq)
    best = [1] * n
    prev = [-1] * n
    for i in range(n):
        for j in range(i):
            if seq[j][0] < seq[i][0] and best[j] + 1 > best[i]:
                best[i], prev[i] = best[j] + 1, j
    i = max(range(n), key=lambda k: (best[k], -k))
    keep: set[int] = set()
    while i != -1:
        keep.add(i)
        i = prev[i]
    return [seq[k][1] for k in range(n) if k not in keep]


def _unique_keys(items: list[str]) -> list[str]:
    seen: dict[str, int] = defaultdict(int)
    out = []
    for key in items:
        seen[key] += 1
        out.append(key if seen[key] == 1 else f"{key}#{seen[key]}")
    return out


def compare_variant(match: VariantMatch, *, tables: bool = True) -> list[Diff]:
    v = match.variant

    def applies(scope: str | None) -> bool:
        return section_applies_to_scope({"scope": scope}, match.report_scope)

    out: list[Diff] = []
    pairs = match.chapter_pairs
    paired_codes = {sc.code for sc in pairs.values()}
    for wc in match.word.chapters:
        sc = pairs.get(wc.index)
        if sc is None:
            out.append(Diff(v, "chapter_missing", wc.title, f"Word 第 {wc.index} 章，{len(wc.sections)} 节"))
            continue
        if sc.scope_norm(match.report_scope) != wc.norm:
            out.append(Diff(v, "chapter_title", sc.code, f"种子「{sc.scope_title(match.report_scope)}」/ Word「{wc.title}」"))
        if not applies(sc.scope):
            out.append(Diff(v, "scope", sc.code, f"章出现在本口径 Word，scope={sc.scope}"))
    for sc in match.chapters:
        if applies(sc.scope) and sc.code not in paired_codes:
            out.append(Diff(v, "chapter_extra", sc.code, f"「{sc.title}」本口径 Word 无此章"))
    order_seq = [(pairs[wc.index].position, pairs[wc.index].code) for wc in match.word.chapters if wc.index in pairs]
    out += [Diff(v, "chapter_order", code) for code in _out_of_order(order_seq)]

    by_ws: dict[int, SeedSection] = {}
    for sc in match.chapters:
        for s in sc.sections:
            hit = match.matches.get(s.position)
            if hit is not None:
                wc, ws = hit
                by_ws[id(ws)] = s
                target = pairs.get(wc.index)
                if target is None or target.code != s.chapter_code:
                    where = f"（对应种子章「{target.code}」）" if target is not None else "（该章种子缺失）"
                    out.append(Diff(v, "misplaced", s.key, f"种子在「{s.chapter_code}」，Word 在第 {wc.index} 章「{wc.title}」{where}"))
                if not applies(s.scope):
                    out.append(Diff(v, "scope", s.key, f"出现在本口径 Word，scope={s.scope}"))
                if s.norm != ws.norm:
                    out.append(Diff(v, "section_title", s.key, f"种子「{s.title[:50]}」/ Word「{ws.title[:50]}」"))
            elif s.position not in match.body_leaves and applies(s.scope):
                sub = "guidance" if s.title.strip().startswith("【") else "other"
                out.append(Diff(v, "json_only", s.key, f"{sub}：{s.title[:60]}"))

    keys = _unique_keys([f"{wc.title}/{ws.title}" for wc, ws in match.word_unmatched])
    for key, (wc, ws) in zip(keys, match.word_unmatched):
        out.append(Diff(v, "word_only", key, f"第 {wc.index} 章，H{ws.level}，{len(ws.tables)} 张表"))

    for wc in match.word.chapters:
        sc = pairs.get(wc.index)
        if sc is None:
            continue
        seq = [
            (by_ws[id(ws)].position, by_ws[id(ws)].key)
            for ws in wc.sections
            if id(ws) in by_ws and by_ws[id(ws)].chapter_code == sc.code
        ]
        out += [Diff(v, "order", key, f"第 {wc.index} 章「{wc.title}」") for key in _out_of_order(seq)]

    if tables:
        out += _table_diffs(match)
    return sorted(set(out))


















def _table_diffs(match: VariantMatch) -> list[Diff]:
    by_pos = match.sections_by_position()
    out: list[Diff] = []
    for pos, (_wc, ws) in match.matches.items():
        s = by_pos[pos]
        out += [Diff(match.variant, *d) for d in compare_tables(s.key, s.raw.get("tables") or [], ws.tables)]
    for pos, wc in match.body_leaves.items():
        s = by_pos[pos]
        out += [Diff(match.variant, *d) for d in compare_tables(s.key, s.raw.get("tables") or [], wc.tables)]
    return out


# ── 种子自身一致性 + 报告 + 白名单 ───────────────────────────────────────────


def seed_diffs(std: str, chapters: list[SeedChapter]) -> list[Diff]:
    """与 Word 无关的种子问题：缺 scope、父节点解析不到、同号、同 section_id。"""
    tag = f"{std}_seed"
    out: list[Diff] = []
    ids: dict[str, int] = defaultdict(int)
    for sc in chapters:
        ids[sc.section_id or ""] += 1
        for s in sc.sections:
            ids[s.section_id or ""] += 1
            if s.scope is None:
                out.append(Diff(tag, "seed_scope_missing", s.key, s.title[:60]))
            if not s.parent_ok:
                out.append(Diff(tag, "seed_parent_missing", s.key, str(s.raw.get("parent_section_id"))))
            if s.key != s.code:
                out.append(Diff(tag, "seed_duplicate_code", s.key, s.title[:60]))
    for sid, n in ids.items():
        if sid and n > 1:
            out.append(Diff(tag, "seed_duplicate_id", sid, f"{n} 次"))
    return out


_WORD_CACHE: dict[str, WordDoc] = {}


def cached_word_variant(variant: str) -> WordDoc:
    if variant not in _WORD_CACHE:
        _WORD_CACHE[variant] = load_word_variant(variant)
    return _WORD_CACHE[variant]


def match_all(seeds: dict[str, dict[str, Any]] | None = None,
              variants: tuple[str, ...] = VARIANT_KEYS) -> dict[str, VariantMatch]:
    """四个口径的匹配结果；``seeds`` 可传入内存中的种子 dict（修复器预演用）。"""
    loaded: dict[str, list[SeedChapter]] = {}
    out: dict[str, VariantMatch] = {}
    for variant in variants:
        std = variant.split("_", 1)[0]
        if std not in loaded:
            loaded[std] = load_seed((seeds or {}).get(std) or std)
        out[variant] = match_variant(cached_word_variant(variant), loaded[std], variant)
    return out


def build_report(seeds: dict[str, dict[str, Any]] | None = None,
                 variants: tuple[str, ...] = VARIANT_KEYS, *, tables: bool = True) -> list[Diff]:
    matches = match_all(seeds, variants)
    diffs: list[Diff] = []
    for m in matches.values():
        diffs += compare_variant(m, tables=tables)
    for std in sorted({v.split("_", 1)[0] for v in variants}):
        diffs += seed_diffs(std, next(m.chapters for v, m in matches.items() if v.startswith(std)))
    return sorted(set(diffs))


def check_against_allowlist(diffs: list[Diff], allowlist: list[dict[str, Any]]) -> tuple[list[Diff], list[dict]]:
    """返回（白名单外差异, 已失效白名单条目）。条目缺理由直接报错。"""
    for entry in allowlist:
        if not str(entry.get("reason") or "").strip():
            raise ValueError(f"白名单条目缺理由：{entry}")
    allowed = {(e["variant"], e["kind"], e["key"]) for e in allowlist}
    present = {(d.variant, d.kind, d.key) for d in diffs}
    unexpected = [d for d in diffs if (d.variant, d.kind, d.key) not in allowed]
    stale = [e for e in allowlist if (e["variant"], e["kind"], e["key"]) not in present]
    return unexpected, stale

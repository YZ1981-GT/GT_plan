"""上年附注 Word → 本年附注表格期初列：解析、逐表规划、按确认写入（纯函数，不读写库）。

spec: consol-note-header-library-and-node-reaggregate R6.6（design §15.6）。服务层 ``note_prior_year_import_service``
装载本年附注、落库、盖「最近一次编辑」。

- 解析：``note_word_table.iter_blocks``（下钻内容控件：平台导出的交付件每节都包在内容控件里）+ ``parse_word_table``
  （网格、表头行数、表头空白格归一、逻辑列、数据行）；标题层级与上传解析同一读法（``docx_paragraph_heading`` +
  ``nest_by_numbering``）。表名候选 = 表前（上一张表 / 标题之后）的全部非空段落（表名与表之间常隔着说明文字）。
- 对节：本年章节标题 ↔ 上年标题（``normalize_title``，同名按出现序号对应）；上年一节的表 = 该标题到下一个同级或更高级
  标题之间的全部表（含小标题下的表）。
- 对表：本年表的期初一侧金额列（期初 / 上期，不含变动列）↔ 上年表期末一侧同计量列（本年列期间平移到期末 / 本期后
  走标题行库 ``match_columns``；上年表头没写期间时按它的「（续：期初数）」列续表补对侧期间，与合并选表同一规则
  ``fill_main_period``）。表名（去期间括注与续表标记）命中 → 选它；没有表名证据时：上年该节只有一张表且本年
  只有一张要导入的表 / 两年表数相同且同位置的表列全对上 → 选它；否则按列结构择优、要人工确认。
  期初列全部对上且选表有名字或位置证据 ⇒ ``auto``；否则 ``needs_confirm`` / ``no_source`` 并写原因。
- 行按「行名 + 同名序号」对应；只填空格（不覆盖已有值）。
"""
from __future__ import annotations

import io
from collections import Counter, defaultdict
from dataclasses import dataclass, field
from typing import Any, Iterable, Sequence

from app.services.note_section_kind import normalize_title
from app.services.note_table_header_library import (
    ColumnKey,
    base_key,
    column_keys,
    continuation_links,
    is_continuation_name,
    load_vocabulary,
    movement_kind,
    name_period,
    norm,
    path_keys,
    period_side,
    table_key,
)
from app.services.note_table_header_match import fill_main_period, match_columns
from app.services.note_word_table import _PLACEHOLDER_ROW, WordTable, iter_blocks, parse_word_table

__all__ = [
    "STATUS_AUTO",
    "STATUS_CONFIRM",
    "STATUS_NONE",
    "PriorDoc",
    "PriorHeading",
    "PriorTable",
    "TablePlan",
    "apply_col_map",
    "match_sections",
    "note_tables",
    "opening_columns",
    "parse_amount",
    "parse_prior_docx",
    "plan_note_tables",
]

STATUS_AUTO, STATUS_CONFIRM, STATUS_NONE = "auto", "needs_confirm", "no_source"
#: 本年期初一侧期间 → 上年期末一侧期间
_SHIFT = {"prior": "end", "prev": "cur"}
#: 有这些选表依据且期初列全部对上才自动写入
_AUTO_METHODS = frozenset({"name", "single_table", "position"})


#: 表名候选里认作「表名」的最长字数（更长的是说明文字：平台导出的表前常是一段正文）
_CAPTION_MAX = 40


@dataclass
class PriorTable:
    id: str
    heading: int | None  # 所在标题（``PriorDoc.headings`` 下标），文档开头无标题时为 None
    captions: list[str]
    word: WordTable
    keys: list[ColumnKey]
    #: 列续表（「（续：期初数）」接在后面、同一批行补期间）：主表表头没写期间时按续表补对侧期间
    continuations: list[str] = field(default_factory=list)

    @property
    def caption(self) -> str:
        """展示用表名：离表最近的短段落；表前只有长段说明时取它首行截断。"""
        short = next((c for c in reversed(self.captions) if len(c) <= _CAPTION_MAX and "\n" not in c), None)
        if short is not None:
            return short
        first = self.captions[-1].split("\n")[0] if self.captions else ""
        return first if len(first) <= _CAPTION_MAX else f"{first[:_CAPTION_MAX]}…"

    def columns(self) -> list[dict[str, Any]]:
        return [{"col": k.index, "label": k.label} for k in self.keys[1:]]


@dataclass
class PriorHeading:
    index: int
    level: int
    title: str
    norm: str
    text: str = ""  # 标题之后、下一个标题之前的段落（上年正文）
    tables: list[str] = field(default_factory=list)  # 子树里的表（含小标题下的表）
    parent: int | None = None


@dataclass
class PriorDoc:
    headings: list[PriorHeading]
    tables: dict[str, PriorTable]


@dataclass
class TablePlan:
    key: str
    note_id: str
    note_section: str
    section_title: str
    table_index: int
    table_name: str
    opening: list[ColumnKey]
    status: str
    reason: str | None = None
    prior_table_id: str | None = None
    method: str | None = None
    col_map: dict[int, int] = field(default_factory=dict)
    candidates: list[str] = field(default_factory=list)
    candidate_maps: dict[str, dict[int, int]] = field(default_factory=dict)

    def to_dict(self) -> dict[str, Any]:
        return {
            "key": self.key, "note_id": self.note_id, "note_section": self.note_section,
            "section_title": self.section_title, "table_index": self.table_index, "table_name": self.table_name,
            "opening_columns": [{"col": k.index, "label": k.label} for k in self.opening],
            "status": self.status, "reason": self.reason, "method": self.method,
            "prior_table_id": self.prior_table_id, "col_map": {str(c): p for c, p in self.col_map.items()},
            "candidates": self.candidates,
            "candidate_maps": {t: {str(c): p for c, p in m.items()} for t, m in self.candidate_maps.items()},
        }


# ─── 解析 ────────────────────────────────────────────────────────────────────


def parse_prior_docx(data: bytes) -> PriorDoc:
    """上年附注 .docx → 标题树 + 表格（文档顺序）。"""
    import docx
    from docx.text.paragraph import Paragraph

    from app.services.note_document_parser import docx_paragraph_heading, nest_by_numbering
    from app.services.note_word_tree import _MARKER

    document = docx.Document(io.BytesIO(data))
    events: list[tuple[str, Any]] = []  # ("h", [level, text, outline_only]) | ("p", text) | ("t", (captions, el))
    style_levels: dict[str, int | None] = {}
    pending: list[str] = []
    for tag, el in iter_blocks(document.element.body):
        if tag == "tbl":
            events.append(("t", (pending, el)))
            pending = []
            continue
        paragraph = Paragraph(el, document)
        text = _MARKER.sub("", paragraph.text or "").strip()
        if not text:
            continue
        level, from_style = docx_paragraph_heading(paragraph, el, style_levels)
        events.append(("h", [level, text, not from_style]) if level is not None else ("p", text))
        pending = [text] if level is not None else [*pending, text]
    heads = [e for kind, e in events if kind == "h"]
    for head, level in zip(heads, nest_by_numbering([h[0] for h in heads], [h[1] for h in heads], [h[2] for h in heads])):
        head[0] = level

    headings: list[PriorHeading] = []
    tables: dict[str, PriorTable] = {}
    stack: list[PriorHeading] = []  # 当前标题链（层级递增）
    for kind, payload in events:
        if kind == "h":
            level, title, _outline = payload
            while stack and stack[-1].level >= level:
                stack.pop()
            head = PriorHeading(len(headings), level, title, normalize_title(title),
                                parent=stack[-1].index if stack else None)
            headings.append(head)
            stack.append(head)
        elif kind == "p":
            if stack:
                stack[-1].text = f"{stack[-1].text}\n{payload}" if stack[-1].text else payload
        else:
            captions, el = payload
            word = parse_word_table(el, captions[-1] if captions else "")
            tid = f"t{len(tables) + 1}"
            # 表头没写期间时用表名期间（「按账龄披露（期末数）」）：取最近一个带期间括注的表名候选
            name = next((c for c in reversed(captions) if name_period(c)), "")
            keys = path_keys(word.header_paths, name=name)
            tables[tid] = PriorTable(tid, stack[-1].index if stack else None, list(captions), word, keys)
            for head in stack:
                head.tables.append(tid)
    _link_continuations(list(tables.values()))
    return PriorDoc(headings, tables)


def _link_continuations(tables: list[PriorTable]) -> None:
    """续表认主表（标题行库 ``continuation_links``：同基名 / 紧挨着的上一张，行名更吻合者）；只记同一批行补期间的
    列续表。只凭「紧挨着」认的须在同一标题下（交付件删掉了空主表时，续表不认到上一节的表上）。"""
    def link_name(t: PriorTable) -> str:
        return next((c for c in reversed(t.captions) if is_continuation_name(c) or name_period(c)), t.caption)

    shells = [{"name": link_name(t), "rows": [{"label": r[0] if r else ""} for r in t.word.rows]} for t in tables]
    for i, link in continuation_links(shells).items():
        main, cont = (tables[link.of] if link.of is not None else None), tables[i]
        if main is not None and link.kind == "columns" and (link.via != "previous" or main.heading == cont.heading):
            main.continuations.append(cont.id)


def match_sections(titles: Sequence[str], doc: PriorDoc) -> list[int | None]:
    """本年各章节（按附注顺序）对应的上年标题：标题规范化后相同，同名按出现序号对应。"""
    by_norm: dict[str, list[int]] = defaultdict(list)
    for head in doc.headings:
        if head.norm:
            by_norm[head.norm].append(head.index)
    seen: Counter[str] = Counter()
    out: list[int | None] = []
    for title in titles:
        key = normalize_title(title)
        nth, seen[key] = seen[key], seen[key] + 1
        hits = by_norm.get(key) or []
        out.append(hits[nth] if key and nth < len(hits) else None)
    return out


# ─── 规划 ────────────────────────────────────────────────────────────────────


def note_tables(table_data: Any) -> list[dict]:
    """附注的表：多表 ``_tables``，否则顶层即唯一一张表。"""
    if not isinstance(table_data, dict) or not table_data:
        return []
    tables = table_data.get("_tables")
    if isinstance(tables, list):
        return [t for t in tables if isinstance(t, dict)]
    return [table_data] if isinstance(table_data.get("rows"), list) else []


def opening_columns(keys: Sequence[ColumnKey]) -> list[ColumnKey]:
    """期初一侧的金额列（期初余额 / 上年年末余额 / 上期发生额…，不含变动列）：从上年期末一侧取数。"""
    vocab = load_vocabulary()
    return [k for k in keys[1:]
            if k.kind == "amount" and period_side(k.period) == "previous" and movement_kind(k, vocab) is None]


def _target_keys(keys: Sequence[ColumnKey], opening: Sequence[ColumnKey]) -> list[ColumnKey]:
    """本年表的对应键：期初一侧列的期间平移到期末一侧（期初 → 期末、上期 → 本期），其余列不参与对应。"""
    wanted = {k.index for k in opening}
    out = [keys[0]]
    for k in keys[1:]:
        if k.index in wanted:
            out.append(ColumnKey(k.index, k.path, _SHIFT[k.period], k.measure, k.kind, k.source))
        else:
            out.append(ColumnKey(k.index, k.path, None, (), "skip", k.source))
    return out


def _col_map(target: Sequence[ColumnKey], prior: PriorTable, doc: PriorDoc) -> tuple[dict[int, int], int]:
    """本年对应键 → 上年表列。上年表头没写期间（多见于有「（续：期初数）」续表的主表）时按续表补对侧期间，
    补不出再取本年对应键的唯一期间（与合并选表同一规则 ``fill_main_period``）。"""
    keys = fill_main_period(prior.keys, [doc.tables[c].keys for c in prior.continuations], target)
    match = match_columns(target, [keys])
    return {c: src for c, (_pos, src) in match.col_map.items()}, match.strong_count


def _name_hit(captions: Iterable[str], want: str) -> bool:
    key, base = table_key(want), base_key(want)
    return any(table_key(c) == key or (base and base_key(c) == base) for c in captions)


def _choose(plan: TablePlan, target: list[ColumnKey], doc: PriorDoc, cands: list[str],
            position: int, note_table_count: int, importing: int) -> None:
    scores: dict[str, tuple[int, int]] = {}
    for tid in cands:
        cmap, strong = _col_map(target, doc.tables[tid], doc)
        plan.candidate_maps[tid] = cmap
        scores[tid] = (strong, len(cmap))
    want = plan.table_name
    hits = [tid for tid in cands if want and _name_hit(doc.tables[tid].captions, want)]
    pool, method = (hits, "name") if hits else (cands, None)
    ranked = sorted(pool, key=lambda t: scores[t], reverse=True)
    unique_top = ranked and (len(ranked) == 1 or scores[ranked[0]] > scores[ranked[1]])
    choice: str | None = None
    if hits:
        if unique_top:
            choice = ranked[0]
        else:
            plan.reason = f"上年该节有 {len(hits)} 张与「{want}」同名的表，表头也分不出取哪一张"
    elif len(cands) == 1 and importing == 1:
        choice, method = cands[0], "single_table"
    elif (len(cands) == note_table_count and len(plan.candidate_maps[cands[position]]) == len(plan.opening)
          and scores[cands[position]] == scores[ranked[0]]):
        choice, method = cands[position], "position"
    elif unique_top and scores[ranked[0]][0] >= 1:
        choice, method = ranked[0], "structure"
        plan.reason = f"上年该节没有同名表，按表头选了「{doc.tables[choice].caption or choice}」，请确认"
    else:
        plan.reason = f"上年该节有 {len(cands)} 张表，没有与「{want or '本表'}」对应的表"
    plan.prior_table_id, plan.method = choice, method
    if choice is None:
        plan.status = STATUS_CONFIRM
        return
    plan.col_map = dict(plan.candidate_maps[choice])
    missing = [k.label for k in plan.opening if k.index not in plan.col_map]
    if missing:
        plan.status = STATUS_CONFIRM
        plan.reason = (f"期初列「{'、'.join(missing)}」在上年表「{doc.tables[choice].caption or choice}」里找不到对应列"
                       + (f"；{plan.reason}" if plan.reason else ""))
    elif method in _AUTO_METHODS:
        plan.status = STATUS_AUTO
    else:
        plan.status = STATUS_CONFIRM


def plan_note_tables(note: dict[str, Any], heading: PriorHeading | None, doc: PriorDoc) -> list[TablePlan]:
    """本年一节里有期初列的表逐张规划。``note``：``{id, note_section, section_title, table_data}``。"""
    tables = note_tables(note.get("table_data"))
    keyed = [(i, t, column_keys(t)) for i, t in enumerate(tables)]
    todo = [(i, t, keys, opening_columns(keys)) for i, t, keys in keyed]
    todo = [x for x in todo if x[3]]
    out: list[TablePlan] = []
    for i, table, keys, opening in todo:
        plan = TablePlan(
            key=f"{note['id']}:{i}", note_id=str(note["id"]), note_section=str(note.get("note_section") or ""),
            section_title=str(note.get("section_title") or ""), table_index=i,
            table_name=str(table.get("name") or ""), opening=opening, status=STATUS_NONE,
        )
        if heading is None:
            plan.reason = f"上年附注里没有找到「{plan.section_title}」这一节"
        elif not heading.tables:
            plan.reason = f"上年附注「{heading.title}」这一节没有表格"
        else:
            plan.candidates = list(heading.tables)
            _choose(plan, _target_keys(keys, opening), doc, plan.candidates, i, len(tables), len(todo))
        out.append(plan)
    return out


# ─── 写入 ────────────────────────────────────────────────────────────────────


def parse_amount(text: Any) -> float | None:
    """Word 表格里的金额文字 → 数值：去千分位，（）/ () 表负数，空 / - / — 视为无值。"""
    s = str(text or "").replace(",", "").replace("，", "").strip()
    negative = s[:1] in ("(", "（") and s[-1:] in (")", "）")
    s = s.strip("()（）").strip()
    if s in ("", "-", "—", "--"):
        return None
    try:
        value = float(s)
    except ValueError:
        return None
    return -value if negative else value


def _row_key(label: Any) -> str:
    key = norm(label).rstrip(":")
    return "" if not key or _PLACEHOLDER_ROW.search(key) else key


def _keyed(labels: Iterable[Any]) -> list[tuple[str, int] | None]:
    seen: Counter[str] = Counter()
    out: list[tuple[str, int] | None] = []
    for label in labels:
        key = _row_key(label)
        if key:
            out.append((key, seen[key]))
            seen[key] += 1
        else:
            out.append(None)
    return out


def apply_col_map(table: dict, prior: WordTable, col_map: dict[int, int]) -> int:
    """按 ``col_map``（本年列 → 上年列，均为表头列号，第 0 列是行名）把上年值填进本年表；返回填入格数。

    行按「行名 + 同名序号」对应；只填空格（已有值不覆盖）。``values`` 不含行名列 ⇒ 下标 = 列号 − 1。
    """
    source = {key: row for key, row in zip(_keyed(r[0] if r else "" for r in prior.rows), prior.rows) if key}
    rows = [r for r in table.get("rows") or [] if isinstance(r, dict) and isinstance(r.get("values"), list)]
    filled = 0
    for key, row in zip(_keyed(r.get("label") for r in rows), rows):
        src = source.get(key) if key else None
        if src is None:
            continue
        values = row["values"]
        for cur, prev in col_map.items():
            value = parse_amount(src[prev]) if 0 < prev < len(src) else None
            if value is None:
                continue
            while len(values) < cur:
                values.append(None)
            if values[cur - 1] in (None, ""):
                values[cur - 1] = value
                filled += 1
    return filled

"""
ContentExtractor: 文档文本提取服务
从 PDF/DOCX/XLSX/TXT/MD 等格式文件中提取纯文本内容。
工厂模式按文件扩展名分派提取器，60s 超时。
"""

from __future__ import annotations

import asyncio
from dataclasses import dataclass
from pathlib import Path


@dataclass
class ExtractResult:
    """提取结果"""

    content_text: str | None
    status: str  # 'extracted' / 'unsupported_format' / 'extraction_failed'
    error: str | None


def _extract_pdf(file_path: str) -> str:
    """PDF 提取 (PyMuPDF fitz)"""
    import fitz  # type: ignore[import-untyped]

    doc = fitz.open(file_path)
    pages: list[str] = []
    for page in doc:
        text = page.get_text("text").strip()
        if text:
            pages.append(text)
    doc.close()
    return "\n\n".join(pages)


def _extract_docx(file_path: str) -> str:
    """DOCX 提取 (python-docx)"""
    from docx import Document  # type: ignore[import-untyped]

    doc = Document(file_path)
    parts: list[str] = []

    # 段落
    for para in doc.paragraphs:
        text = para.text.strip()
        if text:
            parts.append(text)

    # 表格
    for table in doc.tables:
        for row in table.rows:
            cells = [cell.text.strip() for cell in row.cells if cell.text.strip()]
            if cells:
                parts.append("\t".join(cells))

    return "\n\n".join(parts)


def _extract_xlsx(file_path: str) -> str:
    """XLSX 提取 (openpyxl): 非空 sheet 带名称前缀"""
    from openpyxl import load_workbook

    wb = load_workbook(file_path, read_only=True, data_only=True)
    sheet_texts: list[str] = []

    for ws in wb.worksheets:
        rows: list[str] = []
        for row in ws.iter_rows(values_only=True):
            cells = [str(c).strip() for c in row if c is not None and str(c).strip()]
            if cells:
                rows.append("\t".join(cells))
        if rows:
            sheet_texts.append(f"[{ws.title}]\n" + "\n".join(rows))

    wb.close()
    return "\n\n".join(sheet_texts)


def _extract_text(file_path: str) -> str:
    """TXT/MD 提取: 直接 UTF-8 读取"""
    return Path(file_path).read_text("utf-8")


class ContentExtractor:
    """工厂模式按文件扩展名分派提取器"""

    EXTRACTORS: dict[str, callable] = {
        ".pdf": _extract_pdf,
        ".docx": _extract_docx,
        ".doc": _extract_docx,
        ".xlsx": _extract_xlsx,
        ".xls": _extract_xlsx,
        ".txt": _extract_text,
        ".md": _extract_text,
    }

    @classmethod
    async def extract(cls, file_path: str, timeout: int = 60) -> ExtractResult:
        """提取文本，超时 60s，返回 ExtractResult"""
        ext = Path(file_path).suffix.lower()
        extractor = cls.EXTRACTORS.get(ext)
        if not extractor:
            return ExtractResult(None, "unsupported_format", None)
        try:
            text = await asyncio.wait_for(
                asyncio.to_thread(extractor, file_path), timeout=timeout
            )
            return ExtractResult(text, "extracted", None)
        except asyncio.TimeoutError:
            return ExtractResult(None, "extraction_failed", "timeout_60s")
        except Exception as e:
            return ExtractResult(None, "extraction_failed", str(e))

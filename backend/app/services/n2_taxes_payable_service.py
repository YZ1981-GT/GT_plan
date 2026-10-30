"""N2 应交税费 — 业务逻辑服务

纯函数多税种测算引擎 + 增值税测算引擎 + xlsx导入导出 + TB取数/回写 + 跨底稿合计。

提供功能：
- 负债类TB取数（tb_balance 期末余额，direction=贷）
- 增值税测算（销项-进项）
- 多税种测算（城建税/教育费附加/房产税/土地增值税）
- 出口退税核对
- 导出模板 / 导出数据 / 导入数据
- 审定数回写 trial_balance
- 跨底稿合计（N4税金及附加联动）

科目2221应交税费（贷方/负债类）：期末=期初+贷方-借方
N税费循环最复杂底稿（18 sheet/~180+公式）。

Requirements: 4~8, 12.1-12.4
"""

from __future__ import annotations

import io
import json
import logging
from math import isfinite
from typing import Any
from uuid import uuid4

from openpyxl import Workbook, load_workbook
from openpyxl.styles import Alignment, Font, PatternFill

from app.services.n2_detail_model import (
    N2_DETAIL_FIELDS,
    N2_DETAIL_HEADERS,
    N2_DETAIL_IDENTITY_KEY,
    N2_DETAIL_ITEM_ID,
    N2_DETAIL_STORAGE_FIELD,
    enrich_n2_detail_row,
    persisted_n2_detail_row,
)

logger = logging.getLogger(__name__)

_N2_ACCOUNT_CODE = "2221"


# ═══════════════════════════════════════════════════════════════════════════════
# 增值税测算引擎（纯函数）
# ═══════════════════════════════════════════════════════════════════════════════


def calc_output_vat(sales_amount: float, tax_rate: float) -> float:
    """销项税额 = 销售额 × 适用税率

    Requirements: 4.2
    """
    return sales_amount * tax_rate


def calc_payable_vat(
    output_vat: float, input_vat: float, input_transfer_out: float
) -> float:
    """应交增值税 = 销项税额 - (进项税额 - 进项转出)

    Requirements: 4.2
    """
    return output_vat - (input_vat - input_transfer_out)


def calc_vat_burden_rate(payable_vat: float, sales_amount: float) -> float:
    """增值税税负率 = 应交增值税 / 销售额

    sales_amount=0 时返回 0.0（避免除零）。

    Requirements: 4.6
    """
    if sales_amount == 0:
        return 0.0
    return payable_vat / sales_amount


# ═══════════════════════════════════════════════════════════════════════════════
# 多税种测算引擎（纯函数）
# ═══════════════════════════════════════════════════════════════════════════════


def calc_surtax(vat: float, consumption_tax: float, rate: float) -> float:
    """城建税及附加 = (增值税 + 消费税) × 税率

    rate:
    - 城建税：市区7%(0.07) / 县城5%(0.05) / 其他1%(0.01)
    - 教育费附加：3%(0.03)
    - 地方教育附加：2%(0.02)

    Requirements: 5.2
    """
    return (vat + consumption_tax) * rate


def calc_property_tax_by_value(original_value: float, deduct_rate: float) -> float:
    """房产税从价 = 房产原值 × (1 - 扣除比例) × 1.2%

    deduct_rate 通常为 10%~30% (0.1~0.3)。

    Requirements: 6.2
    """
    return original_value * (1 - deduct_rate) * 0.012


def calc_property_tax_by_rent(rent_income: float) -> float:
    """房产税从租 = 租金收入 × 12%

    Requirements: 6.2
    """
    return rent_income * 0.12


def calc_land_vat(
    appreciation: float,
    tax_rate: float,
    deduct_items: float,
    quick_deduct_coef: float,
) -> float:
    """土地增值税 = 增值额 × 税率 - 扣除项目 × 速算扣除系数

    四级累进：
    - 增值率≤50%:   税率30%, 速算系数0
    - 50%<增值率≤100%: 税率40%, 速算系数5%
    - 100%<增值率≤200%: 税率50%, 速算系数15%
    - 增值率>200%:  税率60%, 速算系数35%

    Requirements: 7.2, 7.4
    """
    return appreciation * tax_rate - deduct_items * quick_deduct_coef


def calc_appreciation_rate(appreciation: float, deduct_items: float) -> float:
    """土地增值税增值率 = 增值额 / 扣除项目

    deduct_items=0 时返回 0.0（避免除零）。

    Requirements: 7.2
    """
    if deduct_items == 0:
        return 0.0
    return appreciation / deduct_items


def determine_lvt_bracket(appreciation_rate: float) -> dict[str, float]:
    """根据增值率确定土地增值税税率档次

    Returns:
        {"tax_rate": float, "quick_deduct_coef": float, "bracket": str}

    Requirements: 7.3
    """
    if appreciation_rate <= 0.5:
        return {"tax_rate": 0.30, "quick_deduct_coef": 0.0, "bracket": "≤50%"}
    elif appreciation_rate <= 1.0:
        return {"tax_rate": 0.40, "quick_deduct_coef": 0.05, "bracket": "50%~100%"}
    elif appreciation_rate <= 2.0:
        return {"tax_rate": 0.50, "quick_deduct_coef": 0.15, "bracket": "100%~200%"}
    else:
        return {"tax_rate": 0.60, "quick_deduct_coef": 0.35, "bracket": ">200%"}


# ═══════════════════════════════════════════════════════════════════════════════
# 出口退税核对（纯函数）
# ═══════════════════════════════════════════════════════════════════════════════


def calc_export_refund(export_sales: float, refund_rate: float) -> float:
    """免抵退税额 = 出口销售额 × 退税率

    Requirements: 8.2
    """
    return export_sales * refund_rate


def calc_export_refund_diff(calculated: float, approved: float) -> float:
    """出口退税差异 = 测算免抵退税额 - 批复免抵退税额

    正数表示测算大于批复，需关注。

    Requirements: 8.3
    """
    return calculated - approved


# ═══════════════════════════════════════════════════════════════════════════════
# 负债类公式（纯函数）
# ═══════════════════════════════════════════════════════════════════════════════


def calc_liability_end_balance(begin: float, credit: float, debit: float) -> float:
    """负债类期末余额 = 期初 + 贷方 - 借方（2221贷方科目）

    Requirements: 12.1
    """
    return begin + credit - debit


def calc_audited_amount(unadjusted: float, aje: float, rje: float) -> float:
    """审定数 = 未审数 + AJE + RJE

    Requirements: 2.3
    """
    return unadjusted + aje + rje


def validate_liability_direction(
    begin: float,
    credit: float,
    debit: float,
    reported_end: float | None = None,
) -> dict[str, Any]:
    """校验负债类贷方公式方向

    负债类贷方：期末 = 期初 + 贷方 − 借方

    Requirements: 12.1, 12.4
    """
    expected = calc_liability_end_balance(begin, credit, debit)
    result: dict[str, Any] = {
        "expected": round(expected, 2),
        "formula": "期末=期初+贷方-借方（负债类贷方）",
        "direction": "credit",
        "account_code": _N2_ACCOUNT_CODE,
    }

    if reported_end is not None:
        diff = reported_end - expected
        result["reported"] = reported_end
        result["difference"] = round(diff, 2)
        result["isValid"] = abs(diff) <= 0.01

    return result


# ═══════════════════════════════════════════════════════════════════════════════
# TB 取数 + 回写
# ═══════════════════════════════════════════════════════════════════════════════


async def get_tb_data(project_id: str, db: Any) -> dict[str, Any]:
    """从 tb_balance 取科目2221应交税费余额数据（负债类/贷方）

    Requirements: 12.2
    """
    import sqlalchemy as sa

    from app.models.audit_platform_models import TbBalance
    from app.services.dataset_query import get_active_filter

    result: dict[str, Any] = {
        "account_code": _N2_ACCOUNT_CODE,
        "account_name": "应交税费",
        "direction": "credit",
        "begin_balance": 0.0,
        "debit_amount": 0.0,
        "credit_amount": 0.0,
        "end_balance": 0.0,
        "found": False,
    }

    try:
        active_filter = get_active_filter(project_id)
        stmt = (
            sa.select(
                TbBalance.begin_balance,
                TbBalance.debit_amount,
                TbBalance.credit_amount,
                TbBalance.end_balance,
            )
            .where(
                TbBalance.project_id == str(project_id),
                TbBalance.standard_account_code == _N2_ACCOUNT_CODE,
                active_filter,
            )
            .limit(1)
        )
        row = (await db.execute(stmt)).fetchone()
        if row:
            result["begin_balance"] = _parse_num(row.begin_balance)
            result["debit_amount"] = _parse_num(row.debit_amount)
            result["credit_amount"] = _parse_num(row.credit_amount)
            result["end_balance"] = _parse_num(row.end_balance)
            result["found"] = True
    except Exception as e:  # noqa: BLE001
        logger.warning("N2 service: TB 取数失败 project_id=%s: %s", project_id, e)

    return result


async def writeback_tb(project_id: str, audited_amount: float, db: Any) -> dict[str, Any]:
    """回写审定数到 trial_balance（科目2221，期末余额，贷方/负债类）

    Requirements: 12.3
    """
    import sqlalchemy as sa

    try:
        result = await db.execute(
            sa.text(
                "UPDATE trial_balance "
                "SET audited_amount = :amount "
                "WHERE project_id = :pid AND standard_account_code = :code"
            ),
            {
                "amount": audited_amount,
                "pid": str(project_id),
                "code": _N2_ACCOUNT_CODE,
            },
        )
        await db.flush()
        return {
            "success": True,
            "account_code": _N2_ACCOUNT_CODE,
            "audited_amount": audited_amount,
            "rows_affected": result.rowcount,
        }
    except Exception as e:  # noqa: BLE001
        logger.error("N2 service: TB 回写失败 project_id=%s: %s", project_id, e)
        return {"success": False, "error": str(e)}


# ═══════════════════════════════════════════════════════════════════════════════
# 跨底稿合计（N4联动）
# ═══════════════════════════════════════════════════════════════════════════════


async def get_cross_wp_summary(project_id: str, db: Any) -> dict[str, Any]:
    """获取N2各税种计提合计，供N4税金及附加消费

    从 checklist_responses 中聚合各税种计提数据。

    Requirements: 11.1-11.4
    """
    import sqlalchemy as sa

    tax_accruals: dict[str, float] = {
        "城建税": 0.0,
        "教育费附加": 0.0,
        "地方教育附加": 0.0,
        "房产税": 0.0,
        "土地使用税": 0.0,
        "印花税": 0.0,
        "土地增值税": 0.0,
        "所得税": 0.0,
    }

    try:
        # 查找 N2 审定表中各税种计提额
        # item_id 格式: N2-1-tax-accrual-{taxType}
        rows = (
            await db.execute(
                sa.text(
                    "SELECT cr.item_id, cr.conclusion "
                    "FROM checklist_responses cr "
                    "JOIN working_paper wp ON wp.id = cr.wp_id "
                    "WHERE wp.project_id = :pid "
                    "AND cr.item_id LIKE 'N2-1-tax-accrual-%'"
                ),
                {"pid": str(project_id)},
            )
        ).fetchall()

        for row in rows:
            item_id = row.item_id
            tax_type = item_id.replace("N2-1-tax-accrual-", "")
            amount = _parse_num(row.conclusion)
            if tax_type in tax_accruals:
                tax_accruals[tax_type] = amount

    except Exception as e:  # noqa: BLE001
        logger.warning("N2 service: 跨底稿合计查询失败: %s", e)

    total = sum(tax_accruals.values())
    return {
        "project_id": str(project_id),
        "tax_accruals": tax_accruals,
        "total_accrual": round(total, 2),
        "source": "N2-应交税费",
        "target": "N4-税金及附加",
    }


# ═══════════════════════════════════════════════════════════════════════════════
# 导入导出配置
# ═══════════════════════════════════════════════════════════════════════════════

_ROW_LIMIT = 500

# N2 各sheet导出配置（多税种分sheet导出）
_SHEET_CONFIGS: dict[str, dict[str, Any]] = {
    "N2-2": {
        "title": "N2-2 明细表",
        "headers": N2_DETAIL_HEADERS,
        "fields": N2_DETAIL_FIELDS,
        "item_id": N2_DETAIL_ITEM_ID,
        "storage_field": N2_DETAIL_STORAGE_FIELD,
        "identity_key": N2_DETAIL_IDENTITY_KEY,
        "persist": persisted_n2_detail_row,
        "enrich": enrich_n2_detail_row,
    },
    "N2-6": {
        "title": "N2-6 增值税测算",
        "headers": [
            "期间", "销售额", "销项税额", "进项税额",
            "进项转出", "应交增值税", "已交税额", "未交税额",
        ],
        "fields": [
            "period", "salesAmount", "outputVat", "inputVat",
            "inputTransferOut", "payableVat", "paidVat", "unpaidVat",
        ],
    },
    "N2-8": {
        "title": "N2-8 其他税费测算",
        "headers": [
            "税种", "计税依据", "税率", "应交税额",
            "已交税额", "未交税额", "地区", "政策依据", "备注",
        ],
        "fields": [
            "taxType", "taxBase", "taxRate", "payableAmount",
            "paidAmount", "unpaidAmount", "region", "policyBasis", "remark",
        ],
    },
    "N2-9": {
        "title": "N2-9 房产税测算",
        "headers": [
            "房产名称", "计税方式", "计税依据", "扣除比例",
            "税率", "应交房产税", "备注",
        ],
        "fields": [
            "propertyName", "taxMethod", "taxBase", "deductRate",
            "taxRate", "payableAmount", "remark",
        ],
    },
    "N2-10": {
        "title": "N2-10 土地增值税测算",
        "headers": [
            "项目名称", "转让收入", "扣除项目金额", "增值额",
            "增值率", "适用税率", "速算扣除系数",
        ],
        "fields": [
            "projectName", "transferIncome", "deductItems", "appreciation",
            "appreciationRate", "taxRate", "quickDeductCoef",
        ],
    },
}

_NUMERIC_FIELDS: set[str] = {
    "unadjBegin", "unadjPayable", "unadjPaid", "unadjEnd",
    "beginAdjust", "beginRje", "ajePayable", "ajePaid", "rjePayable", "rjePaid",
    "audBegin", "audPayable", "audPaid", "audEnd",
    "taxBase", "taxRate", "beginBalance", "periodAccrual", "periodPayment",
    "endBalance", "declarationAmount", "difference", "exemptionAmount",
    "salesAmount", "outputVat", "inputVat", "inputTransferOut",
    "payableVat", "paidVat", "unpaidVat",
    "payableAmount", "paidAmount", "unpaidAmount",
    "deductRate", "transferIncome", "deductItems", "appreciation",
    "appreciationRate", "quickDeductCoef",
}


# ═══════════════════════════════════════════════════════════════════════════════
# 导出辅助
# ═══════════════════════════════════════════════════════════════════════════════


def _style_header_row(ws: Any) -> None:
    """给表头行设置样式"""
    header_font = Font(bold=True, size=11)
    header_fill = PatternFill(start_color="D9E1F2", end_color="D9E1F2", fill_type="solid")
    header_align = Alignment(horizontal="center", vertical="center", wrap_text=True)

    for cell in ws[1]:
        cell.font = header_font
        cell.fill = header_fill
        cell.alignment = header_align
        ws.column_dimensions[cell.column_letter].width = max(
            len(str(cell.value or "")) * 2 + 4, 12
        )
    ws.freeze_panes = "A2"


def _parse_num(v: Any) -> float:
    """安全解析数值，None/空/非数值/非有限数均归零。"""
    if v is None or v == "":
        return 0.0
    try:
        value = float(v)
    except (ValueError, TypeError):
        return 0.0
    return value if isfinite(value) else 0.0


def _parse_import_num(v: Any, *, sheet: str, row: int, header: str) -> float:
    """导入边界拒绝 Infinity/NaN；普通非数值文本保持历史归零口径。"""
    if v is None or v == "":
        return 0.0
    try:
        value = float(v)
    except (ValueError, TypeError):
        return 0.0
    if not isfinite(value):
        raise ValueError(f"Sheet '{sheet}' 第 {row} 行「{header}」不是有限数")
    return value


def _safe_str(val: Any) -> str:
    if val is None:
        return ""
    return str(val).strip()


# ═══════════════════════════════════════════════════════════════════════════════
# 导出模板
# ═══════════════════════════════════════════════════════════════════════════════


async def export_template(wp_id: str, db: Any, sheet: str | None = None) -> io.BytesIO:
    """生成空白xlsx模板（多sheet或单sheet）

    Sheet列表: N2-2(A:Q 17列), N2-6(增值税8列), N2-8(其他税费9列),
              N2-9(房产税7列), N2-10(土增税7列)

    Requirements: 12.2
    """
    wb = Workbook()
    wb.remove(wb.active)

    configs_to_export = _SHEET_CONFIGS
    if sheet and sheet in _SHEET_CONFIGS:
        configs_to_export = {sheet: _SHEET_CONFIGS[sheet]}

    for _sheet_key, config in configs_to_export.items():
        ws = wb.create_sheet(title=config["title"][:31])
        for col_idx, header in enumerate(config["headers"], 1):
            ws.cell(row=1, column=col_idx, value=header)
        _style_header_row(ws)

    buffer = io.BytesIO()
    wb.save(buffer)
    buffer.seek(0)
    return buffer


# ═══════════════════════════════════════════════════════════════════════════════
# 导出数据
# ═══════════════════════════════════════════════════════════════════════════════


async def export_data(wp_id: str, db: Any, sheet: str | None = None) -> io.BytesIO:
    """导出当前数据为xlsx（多sheet或单sheet）

    从 checklist_responses 读取已保存的数据并填入xlsx。

    Requirements: 12.2
    """
    import sqlalchemy as sa

    wb = Workbook()
    wb.remove(wb.active)

    configs_to_export = _SHEET_CONFIGS
    if sheet and sheet in _SHEET_CONFIGS:
        configs_to_export = {sheet: _SHEET_CONFIGS[sheet]}

    for sheet_key, config in configs_to_export.items():
        ws = wb.create_sheet(title=config["title"][:31])
        for col_idx, header in enumerate(config["headers"], 1):
            ws.cell(row=1, column=col_idx, value=header)
        _style_header_row(ws)

        item_id = str(config.get("item_id") or f"N2-{sheet_key}-rows")
        storage_field = str(config.get("storage_field") or "remark")
        result = await db.execute(
            sa.text(
                f"SELECT {storage_field} AS payload FROM checklist_responses "
                "WHERE wp_id = :wp_id AND item_id = :item_id LIMIT 1"
            ),
            {"wp_id": wp_id, "item_id": item_id},
        )
        stored = result.fetchone()
        data_rows: list[dict[str, Any]] = []
        payload = getattr(stored, "payload", None) if stored else None
        if payload:
            try:
                parsed = json.loads(payload)
                if isinstance(parsed, list):
                    data_rows = [row for row in parsed if isinstance(row, dict)]
            except (json.JSONDecodeError, TypeError):
                pass

        enrich = config.get("enrich")
        for data_row in data_rows:
            projected = enrich(data_row) if callable(enrich) else data_row
            ws.append(["" if projected.get(field) is None else projected.get(field) for field in config["fields"]])

    buffer = io.BytesIO()
    wb.save(buffer)
    buffer.seek(0)
    return buffer


# ═══════════════════════════════════════════════════════════════════════════════
# 导入数据
# ═══════════════════════════════════════════════════════════════════════════════


async def import_data(
    wp_id: str,
    file_content: bytes,
    db: Any,
    sheet: str | None = None,
    *,
    project_id: Any,
    updated_by: Any,
) -> dict[str, Any]:
    """严格解析 xlsx，并在同一事务中原子覆盖对应 checklist 行。"""
    import sqlalchemy as sa

    if len(file_content) > 10 * 1024 * 1024:
        raise ValueError("文件大小超过 10MB 限制")

    owner_result = await db.execute(
        sa.text(
            "SELECT project_id FROM working_paper "
            "WHERE id = :wp_id AND is_deleted = false"
        ),
        {"wp_id": wp_id},
    )
    actual_project_id = owner_result.scalar_one_or_none()
    if actual_project_id is None:
        raise ValueError(f"底稿不存在或已删除: {wp_id}")
    if str(actual_project_id) != str(project_id):
        raise ValueError("底稿所属项目与已授权项目不一致")

    try:
        wb = load_workbook(io.BytesIO(file_content), read_only=True, data_only=True)
    except Exception as exc:
        raise ValueError("无法解析xlsx文件") from exc

    imported_counts: dict[str, int] = {}
    warnings: list[str] = []
    prepared: list[dict[str, Any]] = []
    try:
        for ws in wb.worksheets:
            ws_title = ws.title.strip()
            matched_key: str | None = None
            matched_config: dict[str, Any] | None = None
            for sheet_key, config in _SHEET_CONFIGS.items():
                if config["title"][:31] == ws_title or sheet_key in ws_title:
                    matched_key = sheet_key
                    matched_config = config
                    break
            if not matched_key or not matched_config or (sheet and matched_key != sheet):
                continue
            if matched_key in imported_counts:
                raise ValueError(f"文件中存在重复的目标 Sheet: {matched_key}")

            header_row = next(ws.iter_rows(min_row=1, max_row=1), ())
            actual_headers = [
                "" if cell.value is None else str(cell.value).strip()
                for cell in header_row
            ]
            while actual_headers and actual_headers[-1] == "":
                actual_headers.pop()
            expected_headers = [str(value) for value in matched_config["headers"]]
            duplicates = sorted({
                value for value in actual_headers
                if value and actual_headers.count(value) > 1
            })
            if actual_headers != expected_headers:
                reason = f"；重复列：{', '.join(duplicates)}" if duplicates else ""
                raise ValueError(
                    f"Sheet '{ws_title}' 表头不匹配：必须严格使用 "
                    f"{len(expected_headers)} 列模板并保持原顺序{reason}"
                )

            parsed_rows: list[dict[str, Any]] = []
            fields = list(matched_config["fields"])
            for row_idx, row in enumerate(
                ws.iter_rows(min_row=2, values_only=True), start=2
            ):
                if row_idx > _ROW_LIMIT + 1:
                    warnings.append(f"Sheet '{ws_title}' 超过 {_ROW_LIMIT} 行限制，已截断")
                    break
                if all(cell is None or str(cell).strip() == "" for cell in row):
                    continue

                identity = str(uuid4())
                identity_key = str(matched_config.get("identity_key") or "rowId")
                row_data: dict[str, Any] = {identity_key: identity}
                for index, (header, field) in enumerate(zip(expected_headers, fields)):
                    raw_value = row[index] if index < len(row) else None
                    row_data[field] = (
                        _parse_import_num(
                            raw_value,
                            sheet=ws_title,
                            row=row_idx,
                            header=header,
                        )
                        if field in _NUMERIC_FIELDS
                        else _safe_str(raw_value)
                    )

                persist = matched_config.get("persist")
                parsed_rows.append(
                    persist(row_data, identity=identity)
                    if callable(persist)
                    else row_data
                )

            item_id = str(matched_config.get("item_id") or f"N2-{matched_key}-rows")
            imported_counts[matched_key] = len(parsed_rows)
            prepared.append({
                "item_id": item_id,
                "storage_field": str(matched_config.get("storage_field") or "remark"),
                "payload": json.dumps(
                    parsed_rows,
                    ensure_ascii=False,
                    allow_nan=False,
                ),
            })
    finally:
        wb.close()

    if not prepared:
        target = f"目标 Sheet {sheet}" if sheet else "可导入的 N2 Sheet"
        raise ValueError(f"文件中未找到{target}")

    # 稳定顺序锁住全部目标行；任一后续写失败由路由统一回滚，文件导入保持原子。
    for target in sorted(prepared, key=lambda item: item["item_id"]):
        await db.execute(
            sa.text(
                "SELECT content_version FROM checklist_responses "
                "WHERE wp_id = :wp_id AND item_id = :item_id FOR UPDATE"
            ),
            {"wp_id": wp_id, "item_id": target["item_id"]},
        )

    for target in prepared:
        storage_field = target["storage_field"]
        await db.execute(
            sa.text(
                "INSERT INTO checklist_responses "
                f"(id, project_id, wp_id, item_id, {storage_field}, updated_by, "
                "created_at, updated_at, content_version) "
                "VALUES (:id, :project_id, :wp_id, :item_id, :payload, :updated_by, "
                "NOW(), NOW(), 1) "
                f"ON CONFLICT (wp_id, item_id) DO UPDATE SET {storage_field} = EXCLUDED.{storage_field}, "
                "updated_by = EXCLUDED.updated_by, updated_at = NOW(), "
                "content_version = COALESCE(checklist_responses.content_version, 1) + 1"
            ),
            {
                "id": str(uuid4()),
                "project_id": str(project_id),
                "wp_id": wp_id,
                "item_id": target["item_id"],
                "payload": target["payload"],
                "updated_by": str(updated_by),
            },
        )

    await db.flush()
    return {
        "imported_count": sum(imported_counts.values()),
        "imported_counts": imported_counts,
        "warnings": warnings,
        "message": "导入成功",
        "_item_ids": [target["item_id"] for target in prepared],
    }

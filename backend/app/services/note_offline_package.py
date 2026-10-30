"""附注离线包身份、签名与导出基线校验。

- ``read_package``：解析上传的离线包，校验 ``_meta_`` 完整性（CI-21）与 1.1 包的 manifest / HMAC 签名；
- ``context_errors``：持锁后按当前项目、年度、准则与章节导出基线复核。
"""
from __future__ import annotations

import hashlib
import hmac
import json
import uuid
from datetime import datetime, timezone
from typing import Any, Iterable

PACKAGE_SCHEMA = "note-offline/1.1"
PACKAGE_VERSION = "1.1"
SIGNING_KEY_ID = "jwt-derived-v1"


def canonical_json(value: Any) -> str:
    return json.dumps(value, ensure_ascii=False, sort_keys=True, separators=(",", ":"), default=str)


def section_baseline_payload(section: dict[str, Any]) -> dict[str, Any]:
    """只覆盖离线导入会写的章节身份、表格与编制指引。"""
    return {
        "section_id": str(section.get("section_id") or section.get("note_section") or ""),
        "note_section": str(section.get("note_section") or ""),
        "section_title": str(section.get("section_title") or ""),
        "guidance_text": section.get("guidance_text") or "",
        "table_data": section.get("table_data") or {},
    }


def section_baseline_digest(section: dict[str, Any]) -> str:
    return hashlib.sha256(canonical_json(section_baseline_payload(section)).encode("utf-8")).hexdigest()


def build_manifest(
    sections: list[dict[str, Any]], *, project_id: uuid.UUID | str, year: int,
    template_type: str, meta_sha256: str,
) -> dict[str, Any]:
    section_ids = [str(s.get("section_id") or s.get("note_section") or "") for s in sections]
    if not section_ids or any(not sid for sid in section_ids) or len(section_ids) != len(set(section_ids)):
        raise ValueError("离线包章节身份为空或重复，拒绝导出")
    return {
        "schema": PACKAGE_SCHEMA,
        "export_id": str(uuid.uuid4()),
        "issued_at": datetime.now(timezone.utc).isoformat(),
        "project_id": str(project_id),
        "year": int(year),
        "template_type": str(template_type),
        "section_ids": section_ids,
        "meta_sha256": meta_sha256,
        "sections": {
            sid: {"baseline_sha256": section_baseline_digest(section)}
            for sid, section in zip(section_ids, sections, strict=True)
        },
    }


def _signing_key() -> bytes:
    from app.core.config import settings

    secret = str(settings.JWT_SECRET_KEY or "")
    if not secret:
        raise RuntimeError("JWT_SECRET_KEY 未配置，无法签发附注离线包")
    return hashlib.sha256(("note-offline-package-v1|" + secret).encode("utf-8")).digest()


def sign_manifest(manifest: dict[str, Any]) -> str:
    return hmac.new(_signing_key(), canonical_json(manifest).encode("utf-8"), hashlib.sha256).hexdigest()


def manifest_signature_valid(manifest: dict[str, Any], signature: str) -> bool:
    try:
        expected = sign_manifest(manifest)
    except Exception:
        return False
    return hmac.compare_digest(expected, str(signature or ""))


def validate_manifest(
    manifest: Any, signature: str, *, section_ids: list[str], meta_sha256: str,
    sheet_section_ids: list[str], key_id: str,
) -> list[str]:
    errors: list[str] = []
    if not isinstance(manifest, dict) or manifest.get("schema") != PACKAGE_SCHEMA:
        return ["离线包 manifest 格式无效"]
    if key_id != SIGNING_KEY_ID:
        errors.append("离线包签名密钥版本不受支持")
    if not manifest_signature_valid(manifest, signature):
        errors.append("离线包签名校验失败，文件身份或基线可能被篡改")
    manifest_ids = manifest.get("section_ids")
    if not isinstance(manifest_ids, list) or any(not isinstance(x, str) or not x for x in manifest_ids):
        errors.append("离线包 manifest.section_ids 格式无效")
        manifest_ids = []
    if len(manifest_ids) != len(set(manifest_ids)):
        errors.append("离线包 manifest.section_ids 含重复章节")
    if manifest_ids != section_ids:
        errors.append("离线包 manifest 与 _meta_.section_ids 不一致")
    if set(sheet_section_ids) != set(section_ids) or len(sheet_section_ids) != len(set(sheet_section_ids)):
        errors.append("离线包章节 sheet 与 _meta_.section_ids 不一致或重复")
    if manifest.get("meta_sha256") != meta_sha256:
        errors.append("离线包 manifest 的 meta 摘要不一致")
    entries = manifest.get("sections")
    if not isinstance(entries, dict) or set(entries) != set(section_ids):
        errors.append("离线包 manifest 章节基线不完整")
    elif any(not isinstance(v, dict) or not v.get("baseline_sha256") for v in entries.values()):
        errors.append("离线包 manifest 章节基线摘要无效")
    try:
        uuid.UUID(str(manifest.get("project_id")))
        int(manifest.get("year"))
    except (TypeError, ValueError):
        errors.append("离线包项目或年度身份无效")
    if manifest.get("template_type") not in {"soe", "listed", "custom"}:
        errors.append("离线包准则身份无效")
    return errors


def context_errors(
    manifest: dict[str, Any], *, project_id: uuid.UUID | str, year: int, template_type: str,
    current_sections: list[dict[str, Any]], write_section_ids: Iterable[str] | None = None,
) -> list[str]:
    """在持有当前章节锁后复核 URL/项目准则和导出基线。"""
    errors: list[str] = []
    if str(manifest.get("project_id")) != str(project_id):
        errors.append("离线包属于其他项目，拒绝导入")
    if manifest.get("year") != int(year):
        errors.append("离线包属于其他年度，拒绝导入")
    if manifest.get("template_type") != str(template_type):
        errors.append("离线包准则与当前项目不一致，请重新导出")
    by_id: dict[str, dict[str, Any]] = {}
    duplicates: set[str] = set()
    for section in current_sections:
        sid = str(section.get("section_id") or section.get("note_section") or "")
        if sid in by_id:
            duplicates.add(sid)
        by_id[sid] = section
    if duplicates:
        errors.append(f"当前项目章节身份重复：{sorted(duplicates)}")
    wanted = set(write_section_ids or manifest.get("section_ids") or [])
    baselines = manifest.get("sections") or {}
    stale: list[str] = []
    for sid in sorted(wanted):
        current = by_id.get(sid)
        expected = (baselines.get(sid) or {}).get("baseline_sha256")
        if current is None or not expected or section_baseline_digest(current) != expected:
            stale.append(sid)
    if stale:
        errors.append(f"以下章节在导出后已变化，请重新导出：{stale}")
    return errors


# ---------------------------------------------------------------------------
# C.0.9：读取并校验上传的离线包（原 note_offline_import_service.validate_import_file 主体）
# ---------------------------------------------------------------------------

_NON_SECTION_SHEETS = ("注意事项", "章节清单", "_meta_")


def read_package(xlsx_bytes: bytes) -> dict[str, Any]:
    """解析上传的离线包并校验，返回 ``ImportValidationResult`` 的构造参数（``errors`` 为空即有效）。

    Checks:
    1. Valid xlsx format
    2. _meta_ sheet exists
    3. _meta_ sheet has section_ids + binding_hash (CI-21)
    4. format_version is compatible：1.0 旧包标 ``legacy_unsigned``（只能预览）；
       1.1 包另校验 manifest、HMAC 签名、密钥版本与章节 sheet 一致
    """
    from io import BytesIO

    from openpyxl import load_workbook

    # 导出服务在模块级引用本模块，meta 解压函数只能在函数内取。
    from app.services.note_offline_export_service import _decompress_meta

    try:
        wb = load_workbook(BytesIO(xlsx_bytes), read_only=True, data_only=True)
    except Exception as e:  # noqa: BLE001 — 任何解析失败都如实回给上传者
        return {"errors": [f"无法解析 xlsx 文件: {e}"]}
    try:
        return _inspect_workbook(wb, _decompress_meta)
    finally:
        wb.close()


def _inspect_workbook(wb: Any, decompress_meta: Any) -> dict[str, Any]:
    errors: list[str] = []

    # Check _meta_ sheet exists
    if "_meta_" not in wb.sheetnames:
        return {"errors": ["缺少 _meta_ 隐藏 sheet，无法导入"]}

    ws_meta = wb["_meta_"]

    # Read meta fields
    meta_data_key = ws_meta.cell(row=1, column=1).value
    meta_data_val = ws_meta.cell(row=1, column=2).value
    section_ids_key = ws_meta.cell(row=2, column=1).value
    section_ids_val = ws_meta.cell(row=2, column=2).value
    binding_hash_key = ws_meta.cell(row=3, column=1).value
    binding_hash_val = ws_meta.cell(row=3, column=2).value
    format_version_val = ws_meta.cell(row=5, column=2).value
    manifest_key = ws_meta.cell(row=6, column=1).value
    manifest_val = ws_meta.cell(row=6, column=2).value
    signature_key = ws_meta.cell(row=7, column=1).value
    signature_val = ws_meta.cell(row=7, column=2).value
    key_id_key = ws_meta.cell(row=8, column=1).value
    key_id_val = ws_meta.cell(row=8, column=2).value

    # Validate meta_data
    if meta_data_key != "meta_data" or not meta_data_val:
        errors.append("_meta_ sheet 缺少 meta_data 字段")

    # Validate section_ids (CI-21)
    section_ids: list[str] = []
    if section_ids_key != "section_ids" or not section_ids_val:
        errors.append("_meta_ sheet 缺少 section_ids（CI-21 校验失败）")
    else:
        try:
            parsed_ids = json.loads(section_ids_val)
            if not isinstance(parsed_ids, list) or any(not isinstance(x, str) or not x for x in parsed_ids):
                raise ValueError
            section_ids = parsed_ids
            if len(section_ids) != len(set(section_ids)):
                errors.append("section_ids 含重复章节")
        except (json.JSONDecodeError, TypeError, ValueError):
            errors.append("section_ids 格式无效")

    # Validate binding_hash (CI-21)
    binding_hash = ""
    if binding_hash_key != "binding_hash" or not binding_hash_val:
        errors.append("_meta_ sheet 缺少 binding_hash（CI-21 校验失败）")
    else:
        binding_hash = str(binding_hash_val)
        if meta_data_val and hashlib.sha256(str(meta_data_val).encode("utf-8")).hexdigest() != binding_hash:
            errors.append("binding_hash 与 meta_data 不一致，文件可能被篡改")

    # 1.0 只允许预览；生产导出的 1.1 还带服务端 HMAC 身份与章节基线。
    format_version = str(format_version_val or "")
    if format_version and format_version not in ("1.0", "1.1"):
        errors.append(f"不支持的格式版本: {format_version}")

    # Decompress meta to verify integrity
    meta_data: dict[str, Any] = {}
    if meta_data_val:
        try:
            meta_data = decompress_meta(str(meta_data_val))
        except Exception as e:  # noqa: BLE001
            errors.append(f"meta_data 解压失败: {e}")

    manifest: dict[str, Any] = {}
    sheet_section_ids: list[str] = []
    for ws in wb.worksheets:
        if ws.title in _NON_SECTION_SHEETS:
            continue
        raw_sid = ws.cell(row=2, column=1).value
        if raw_sid and str(raw_sid).startswith("section_id:"):
            sheet_section_ids.append(str(raw_sid).replace("section_id:", "", 1).strip())
    if meta_data and set(meta_data) != set(section_ids):
        errors.append("meta_data 章节键与 section_ids 不一致")
    if set(sheet_section_ids) != set(section_ids) or len(sheet_section_ids) != len(set(sheet_section_ids)):
        errors.append("章节 sheet 与 section_ids 不一致或重复")
    legacy_unsigned = format_version != "1.1"
    if format_version == "1.1":
        if manifest_key != "package_manifest" or not manifest_val:
            errors.append("离线包缺少 package_manifest")
        if signature_key != "package_signature" or not signature_val:
            errors.append("离线包缺少 package_signature")
        if key_id_key != "signing_key_id" or not key_id_val:
            errors.append("离线包缺少 signing_key_id")
        if manifest_val:
            try:
                parsed_manifest = json.loads(str(manifest_val))
                manifest = parsed_manifest if isinstance(parsed_manifest, dict) else {}
            except (json.JSONDecodeError, TypeError):
                errors.append("package_manifest 格式无效")
        if manifest:
            errors.extend(validate_manifest(
                manifest,
                str(signature_val or ""),
                section_ids=section_ids,
                meta_sha256=binding_hash,
                sheet_section_ids=sheet_section_ids,
                key_id=str(key_id_val or ""),
            ))

    return {
        "errors": errors,
        "section_ids": section_ids,
        "meta_data": meta_data,
        "binding_hash": binding_hash,
        "format_version": format_version,
        "manifest": manifest,
        "legacy_unsigned": legacy_unsigned,
    }

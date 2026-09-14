"""编制说明 ``source_ref`` 的权威来源解析与验证。

本模块只验证调用方提供的引用与冻结模板 lineage，不决定 guidance 是否 exact；
``guidance_inventory`` / service 的接线由上层完成。所有失败均结构化返回，禁止把
不存在、跨模板或摘要漂移降级成“有 kind/path 即有效”。
"""
from __future__ import annotations

import hashlib
import json
import re
from collections import Counter
from dataclasses import asdict, dataclass, replace
from pathlib import Path, PurePosixPath, PureWindowsPath
from typing import Any, Literal, Mapping, Sequence
from xml.etree import ElementTree
from zipfile import BadZipFile, ZipFile

from app.services.workpaper_sync.canonical_paths import (
    REPO_ROOT,
    TEMPLATE_ROOT,
    PathBoundaryError,
    assert_within_root,
    resolve_within_root,
)

SourceRefStatus = Literal["valid", "missing", "invalid", "stale", "cross_template"]

_SHA256_RE = re.compile(r"^[0-9a-fA-F]{64}$")
_MAX_XLSX_SOURCE_CELLS = 200_000
_MAX_EXCEL_ROW = 1_048_576
_MAX_EXCEL_COLUMN = 16_384
_W_NS = "http://schemas.openxmlformats.org/wordprocessingml/2006/main"
_W_BODY = f"{{{_W_NS}}}body"
_W_P = f"{{{_W_NS}}}p"
_W_TBL = f"{{{_W_NS}}}tbl"
_W_T = f"{{{_W_NS}}}t"
_W_BOOKMARK_START = f"{{{_W_NS}}}bookmarkStart"
_W_BOOKMARK_END = f"{{{_W_NS}}}bookmarkEnd"
_W_ID = f"{{{_W_NS}}}id"
_W_NAME = f"{{{_W_NS}}}name"


@dataclass(frozen=True)
class TemplateAuthority:
    """当前 render context 允许引用的一份 canonical/active 模板。"""

    canonical_path: str
    active_path: Path | str | None = None
    wp_codes: tuple[str, ...] = ()
    origin: str = "canonical"
    version: str | None = None
    expected_active_digest: str | None = None


@dataclass(frozen=True)
class SourceRefContext:
    """验证所需的冻结上下文；authority 必须由模板 resolver 提供。"""

    target_wp_code: str
    template_authorities: tuple[TemplateAuthority, ...]
    target_sheet_code: str | None = None
    target_sheet_name: str | None = None
    repo_root: Path = REPO_ROOT
    template_root: Path = TEMPLATE_ROOT


@dataclass(frozen=True)
class SourceRefIssue:
    code: str
    field: str | None
    detail: str


@dataclass(frozen=True)
class SourceRefFacts:
    kind: str | None = None
    normalized_path: str | None = None
    authority_root: str | None = None
    expected_digest: str | None = None
    observed_digest: str | None = None
    document_type: str | None = None
    canonical_template_path: str | None = None
    active_template_origin: str | None = None
    active_template_version: str | None = None
    sheet: str | None = None
    cell_range: str | None = None
    anchor_type: str | None = None
    anchor_match_count: int | None = None
    located_content_digest: str | None = None


@dataclass(frozen=True)
class SourceRefValidation:
    status: SourceRefStatus
    issues: tuple[SourceRefIssue, ...]
    facts: SourceRefFacts

    @property
    def valid(self) -> bool:
        return self.status == "valid"


@dataclass(frozen=True)
class SourceRefBatchValidation:
    results: tuple[SourceRefValidation, ...]
    status_counts: dict[str, int]
    facts_digest: str

    @property
    def all_valid(self) -> bool:
        return bool(self.results) and all(item.valid for item in self.results)


class _RefSchemaError(ValueError):
    def __init__(self, code: str, field: str | None, detail: str) -> None:
        self.code = code
        self.field = field
        self.detail = detail
        super().__init__(detail)


def _failure(
    status: SourceRefStatus,
    code: str,
    detail: str,
    facts: SourceRefFacts,
    *,
    field: str | None = None,
) -> SourceRefValidation:
    return SourceRefValidation(
        status=status,
        issues=(SourceRefIssue(code=code, field=field, detail=detail),),
        facts=facts,
    )


def _stable_digest(value: Any) -> str:
    payload = json.dumps(value, ensure_ascii=False, sort_keys=True, separators=(",", ":"))
    return hashlib.sha256(payload.encode("utf-8")).hexdigest()


def file_sha256(path: Path) -> str:
    """按文件字节计算 SHA-256；mtime/size 只能作缓存键，不能作身份。"""
    digest = hashlib.sha256()
    with path.open("rb") as stream:
        for chunk in iter(lambda: stream.read(1024 * 1024), b""):
            digest.update(chunk)
    return digest.hexdigest()


def _normalize_sha256(raw: Any, *, field: str, required: bool = True) -> str | None:
    if raw is None or (isinstance(raw, str) and not raw.strip()):
        if required:
            raise _RefSchemaError("digest_missing", field, f"{field} 必须是文件 SHA-256")
        return None
    if not isinstance(raw, str) or not _SHA256_RE.fullmatch(raw.strip()):
        raise _RefSchemaError("digest_invalid", field, f"{field} 必须是 64 位十六进制 SHA-256")
    return raw.strip().lower()


def _raw_repo_relative_path(raw: Any, *, field: str = "path") -> str:
    if not isinstance(raw, str) or not raw.strip():
        raise _RefSchemaError("path_missing", field, "source_ref.path 不得为空")
    value = raw.strip()
    if "\x00" in value:
        raise _RefSchemaError("path_boundary_rejected", field, "source_ref.path 含 NUL")
    windows = PureWindowsPath(value)
    normalized = value.replace("\\", "/")
    posix = PurePosixPath(normalized)
    if windows.drive or windows.is_absolute() or posix.is_absolute():
        raise _RefSchemaError("path_not_repo_relative", field, "source_ref.path 必须是仓库相对路径")
    if any(part == ".." for part in posix.parts):
        raise _RefSchemaError("path_boundary_rejected", field, "source_ref.path 不得包含 ..")
    cleaned = posix.as_posix()
    if cleaned in ("", "."):
        raise _RefSchemaError("path_missing", field, "source_ref.path 不得为空")
    return cleaned


def _resolve_repo_path(raw: Any, context: SourceRefContext, *, template_only: bool) -> tuple[Path, str]:
    relative = _raw_repo_relative_path(raw)
    try:
        resolved = resolve_within_root(context.repo_root, relative, boundary="repo")
        if template_only:
            assert_within_root(context.template_root, resolved, boundary="template_root")
    except PathBoundaryError as exc:
        code = "authority_root_not_allowed" if template_only else "path_boundary_rejected"
        raise _RefSchemaError(code, "path", "source_ref.path 不在允许的权威根目录") from exc

    repo_real = Path(context.repo_root).resolve()
    try:
        normalized = resolved.relative_to(repo_real).as_posix()
    except ValueError as exc:
        raise _RefSchemaError(
            "path_boundary_rejected", "path", "source_ref.path 无法规范化为仓库相对路径"
        ) from exc
    return resolved, normalized


def _resolve_active_path(
    authority: TemplateAuthority,
    canonical_path: Path,
    context: SourceRefContext,
) -> Path:
    raw = authority.active_path
    if raw is None:
        return canonical_path
    candidate = Path(raw)
    try:
        if candidate.is_absolute():
            return assert_within_root(context.repo_root, candidate, boundary="repo")
        return resolve_within_root(context.repo_root, str(raw), boundary="repo")
    except PathBoundaryError as exc:
        raise _RefSchemaError(
            "template_authority_invalid",
            "template_authorities",
            "active template path 越出仓库边界",
        ) from exc


def _matching_authority(
    normalized_path: str,
    canonical_path: Path,
    context: SourceRefContext,
) -> tuple[TemplateAuthority, Path] | SourceRefValidation:
    facts = SourceRefFacts(
        normalized_path=normalized_path,
        authority_root="backend/wp_templates",
        canonical_template_path=normalized_path,
    )
    if not context.template_authorities:
        return _failure(
            "invalid",
            "template_authority_missing",
            "验证上下文未提供当前模板 lineage",
            facts,
            field="template_authorities",
        )

    path_matches: list[tuple[TemplateAuthority, Path]] = []
    for authority in context.template_authorities:
        try:
            authority_path, authority_normalized = _resolve_repo_path(
                authority.canonical_path, context, template_only=True
            )
            active_path = _resolve_active_path(authority, authority_path, context)
        except _RefSchemaError as exc:
            return _failure("invalid", exc.code, exc.detail, facts, field=exc.field)
        if authority_normalized == normalized_path:
            path_matches.append((authority, active_path))

    if not path_matches:
        if canonical_path.is_file():
            return _failure(
                "cross_template",
                "template_not_active_for_context",
                "引用指向真实模板，但不属于当前 render/template lineage",
                facts,
                field="path",
            )
        return _failure(
            "missing",
            "source_not_found",
            "引用的模板文件不存在",
            facts,
            field="path",
        )

    code_matches = [
        item
        for item in path_matches
        if not item[0].wp_codes or context.target_wp_code in item[0].wp_codes
    ]
    if not code_matches:
        return _failure(
            "cross_template",
            "template_wp_code_mismatch",
            "模板已登记，但不属于当前 wp_code",
            facts,
            field="path",
        )

    unique_active = {str(item[1].resolve()) for item in code_matches}
    if len(unique_active) != 1:
        return _failure(
            "invalid",
            "template_membership_ambiguous",
            "同一 canonical path 对应多个 active template",
            facts,
            field="template_authorities",
        )
    return code_matches[0]


def _normalized_text(value: str) -> str:
    return re.sub(r"\s+", " ", value or "").strip()


def _validate_local_digest(
    expected: Any,
    observed: str,
    facts: SourceRefFacts,
) -> SourceRefValidation | None:
    if expected is None:
        return None
    try:
        normalized = _normalize_sha256(expected, field="anchor.content_digest")
    except _RefSchemaError as exc:
        return _failure("invalid", exc.code, exc.detail, facts, field=exc.field)
    if normalized != observed:
        return _failure(
            "stale",
            "anchor_content_digest_mismatch",
            "语义锚点内容摘要已变化",
            facts,
            field="anchor.content_digest",
        )
    return None


def _validate_xlsx(
    raw: Mapping[str, Any],
    source_path: Path,
    facts: SourceRefFacts,
) -> SourceRefValidation:
    from openpyxl import load_workbook
    from openpyxl.utils import get_column_letter
    from openpyxl.utils.cell import range_boundaries
    from openpyxl.utils.exceptions import InvalidFileException

    sheet = raw.get("sheet")
    cell_range = raw.get("range")
    if not isinstance(sheet, str) or not sheet.strip():
        return _failure("invalid", "sheet_missing", "xlsx source_ref 必须给出 sheet", facts, field="sheet")
    if not isinstance(cell_range, str) or not cell_range.strip():
        return _failure("invalid", "range_missing", "xlsx source_ref 必须给出 range", facts, field="range")
    sheet = sheet.strip()
    cell_range = cell_range.strip()
    facts = replace(facts, sheet=sheet, cell_range=cell_range)
    if any(token in cell_range for token in ("!", "[", "]", ",")):
        return _failure("invalid", "range_invalid", "range 只能是单一 A1 区域", facts, field="range")
    try:
        min_col, min_row, max_col, max_row = range_boundaries(cell_range)
    except (TypeError, ValueError):
        return _failure("invalid", "range_invalid", "range 不是合法 A1 区域", facts, field="range")
    if not all(isinstance(value, int) for value in (min_col, min_row, max_col, max_row)):
        return _failure("invalid", "range_invalid", "range 边界不完整", facts, field="range")
    if (
        min_col < 1
        or min_row < 1
        or max_col < min_col
        or max_row < min_row
        or max_col > _MAX_EXCEL_COLUMN
        or max_row > _MAX_EXCEL_ROW
    ):
        return _failure("invalid", "range_out_of_bounds", "range 超出 Excel 边界", facts, field="range")
    area = (max_col - min_col + 1) * (max_row - min_row + 1)
    if area > _MAX_XLSX_SOURCE_CELLS:
        return _failure("invalid", "range_too_large", "source_ref.range 超出安全读取上限", facts, field="range")
    normalized_range = (
        f"{get_column_letter(min_col)}{min_row}:"
        f"{get_column_letter(max_col)}{max_row}"
    )
    facts = replace(facts, cell_range=normalized_range)

    workbook = None
    try:
        workbook = load_workbook(source_path, read_only=True, data_only=False, keep_links=False)
        if sheet not in workbook.sheetnames:
            return _failure("stale", "sheet_not_found", "引用的工作表已不存在", facts, field="sheet")
        worksheet = workbook[sheet]
        located: list[dict[str, str]] = []
        for row in worksheet.iter_rows(
            min_row=min_row,
            max_row=max_row,
            min_col=min_col,
            max_col=max_col,
        ):
            for cell in row:
                value = cell.value
                if value is None or (isinstance(value, str) and not value.strip()):
                    continue
                located.append(
                    {
                        "coordinate": cell.coordinate,
                        "data_type": str(cell.data_type or ""),
                        "value": str(value),
                    }
                )
    except (OSError, ValueError, KeyError, BadZipFile, InvalidFileException) as exc:
        return _failure(
            "invalid",
            "source_document_invalid",
            f"xlsx 无法解析：{type(exc).__name__}",
            facts,
            field="path",
        )
    finally:
        if workbook is not None:
            workbook.close()
    if not located:
        return _failure("stale", "range_empty", "引用区域不含任何值或公式", facts, field="range")
    content_digest = _stable_digest(located)
    return SourceRefValidation(
        status="valid",
        issues=(),
        facts=replace(facts, located_content_digest=content_digest),
    )


def _docx_xml(path: Path) -> ElementTree.Element:
    try:
        with ZipFile(path) as package:
            payload = package.read("word/document.xml")
        return ElementTree.fromstring(payload)
    except (OSError, BadZipFile, KeyError, ElementTree.ParseError) as exc:
        raise _RefSchemaError(
            "source_document_invalid", "path", f"docx 无法解析：{type(exc).__name__}"
        ) from exc


def _paragraph_texts(root: ElementTree.Element) -> list[str]:
    return [
        "".join(node.text or "" for node in paragraph.iter(_W_T)).strip()
        for paragraph in root.iter(_W_P)
    ]


def _body_blocks(container: ElementTree.Element):
    """按 body 顺序展平内容控件；遇到 paragraph/table 后不重复遍历其子孙。"""
    for child in list(container):
        if child.tag in (_W_P, _W_TBL):
            yield child
        else:
            yield from _body_blocks(child)


def _parse_anchor(raw: Any) -> dict[str, Any]:
    if isinstance(raw, str):
        value = raw.strip()
        if value == "before_first_table":
            return {"type": value}
        if value.startswith("bookmark:") and value.removeprefix("bookmark:").strip():
            return {"type": "bookmark", "name": value.removeprefix("bookmark:").strip()}
        if value:
            return {"type": "bookmark", "name": value}
        raise _RefSchemaError("anchor_missing", "anchor", "docx anchor 不得为空")
    if not isinstance(raw, Mapping):
        raise _RefSchemaError("anchor_missing", "anchor", "docx source_ref 必须给出语义 anchor")
    anchor = dict(raw)
    anchor_type = str(anchor.get("type") or "").strip()
    if anchor_type not in ("bookmark", "paragraph_text", "before_first_table"):
        raise _RefSchemaError("anchor_schema_invalid", "anchor.type", "不支持的 docx anchor.type")
    anchor["type"] = anchor_type
    return anchor


def _validate_docx(
    raw: Mapping[str, Any],
    source_path: Path,
    facts: SourceRefFacts,
) -> SourceRefValidation:
    try:
        anchor = _parse_anchor(raw.get("anchor"))
        root = _docx_xml(source_path)
    except _RefSchemaError as exc:
        return _failure("invalid", exc.code, exc.detail, facts, field=exc.field)

    anchor_type = anchor["type"]
    facts = replace(facts, anchor_type=anchor_type)
    located_text = ""
    match_count = 0

    if anchor_type == "bookmark":
        name = anchor.get("name")
        if not isinstance(name, str) or not name.strip():
            return _failure("invalid", "anchor_schema_invalid", "bookmark anchor 缺少 name", facts, field="anchor.name")
        name = name.strip()
        elements = list(root.iter())
        starts = [
            (index, element)
            for index, element in enumerate(elements)
            if element.tag == _W_BOOKMARK_START and (element.get(_W_NAME) or "") == name
        ]
        match_count = len(starts)
        facts = replace(facts, anchor_match_count=match_count)
        if not starts:
            return _failure("stale", "anchor_not_found", "bookmark 已不存在", facts, field="anchor.name")
        if len(starts) != 1:
            return _failure("invalid", "anchor_ambiguous", "bookmark name 命中多处", facts, field="anchor.name")
        start_index, start = starts[0]
        bookmark_id = start.get(_W_ID)
        if not bookmark_id:
            return _failure("invalid", "bookmark_id_missing", "bookmarkStart 缺少 w:id", facts, field="anchor.name")
        ends = [
            index
            for index, element in enumerate(elements)
            if element.tag == _W_BOOKMARK_END and element.get(_W_ID) == bookmark_id
        ]
        if not ends:
            return _failure("invalid", "bookmark_unclosed", "bookmarkStart 无匹配 bookmarkEnd", facts, field="anchor.name")
        if len(ends) != 1 or ends[0] <= start_index:
            return _failure("invalid", "anchor_ambiguous", "bookmark end 重复或顺序无效", facts, field="anchor.name")
        located_text = "".join(
            element.text or ""
            for element in elements[start_index + 1 : ends[0]]
            if element.tag == _W_T
        ).strip()
    elif anchor_type == "paragraph_text":
        expected = anchor.get("text")
        mode = str(anchor.get("match") or "exact").strip()
        if not isinstance(expected, str) or not expected.strip() or mode not in ("exact", "normalized"):
            return _failure(
                "invalid",
                "anchor_schema_invalid",
                "paragraph_text 必须给出 text，match 只能是 exact/normalized",
                facts,
                field="anchor",
            )
        expected_value = expected.strip() if mode == "exact" else _normalized_text(expected)
        matches = []
        for text in _paragraph_texts(root):
            observed = text.strip() if mode == "exact" else _normalized_text(text)
            if observed == expected_value:
                matches.append(text)
        match_count = len(matches)
        facts = replace(facts, anchor_match_count=match_count)
        if not matches:
            return _failure("stale", "anchor_not_found", "目标段落已不存在", facts, field="anchor.text")
        if len(matches) != 1:
            return _failure("invalid", "anchor_ambiguous", "目标段落命中多处", facts, field="anchor.text")
        located_text = matches[0]
    else:
        body = next(root.iter(_W_BODY), None)
        if body is None:
            return _failure("invalid", "source_document_invalid", "docx 缺少 w:body", facts, field="path")
        before_table: list[str] = []
        table_found = False
        for block in _body_blocks(body):
            if block.tag == _W_TBL:
                table_found = True
                break
            text = "".join(node.text or "" for node in block.iter(_W_T)).strip()
            if text:
                before_table.append(text)
        facts = replace(facts, anchor_match_count=1 if table_found else 0)
        if not table_found:
            return _failure("stale", "first_table_not_found", "文档已无首个表格边界", facts, field="anchor")
        located_text = "\n".join(before_table)

    normalized = _normalized_text(located_text)
    if not normalized:
        code = "before_first_table_empty" if anchor_type == "before_first_table" else "anchor_empty"
        return _failure("stale", code, "语义锚点未定位到可观察正文", facts, field="anchor")
    content_digest = hashlib.sha256(normalized.encode("utf-8")).hexdigest()
    facts = replace(facts, located_content_digest=content_digest)
    mismatch = _validate_local_digest(anchor.get("content_digest"), content_digest, facts)
    if mismatch is not None:
        return mismatch
    return SourceRefValidation(status="valid", issues=(), facts=facts)


def _validate_template_bound_ref(
    raw: Mapping[str, Any],
    context: SourceRefContext,
    *,
    kind: str,
) -> SourceRefValidation:
    """xlsx/docx 共用：path/authority/digest 门 + 类型专有定位。

    供 ``validate_source_ref`` 与 G-ID registry handler 调用，避免互相递归。
    """
    facts = SourceRefFacts(kind=kind, document_type=kind)
    try:
        expected_digest = _normalize_sha256(raw.get("digest"), field="digest")
        canonical_path, normalized_path = _resolve_repo_path(raw.get("path"), context, template_only=True)
    except _RefSchemaError as exc:
        return _failure("invalid", exc.code, exc.detail, facts, field=exc.field)
    facts = replace(
        facts,
        normalized_path=normalized_path,
        authority_root="backend/wp_templates",
        expected_digest=expected_digest,
        canonical_template_path=normalized_path,
    )

    suffix = canonical_path.suffix.lower()
    expected_suffixes = {"xlsx": {".xlsx", ".xlsm"}, "docx": {".docx"}}
    if suffix not in expected_suffixes[kind]:
        return _failure(
            "invalid",
            "extension_kind_mismatch",
            f"kind={kind} 与文件扩展名 {suffix or '<none>'} 不一致",
            facts,
            field="path",
        )

    authority_match = _matching_authority(normalized_path, canonical_path, context)
    if isinstance(authority_match, SourceRefValidation):
        return replace(
            authority_match,
            facts=replace(
                authority_match.facts,
                kind=kind,
                document_type=kind,
                expected_digest=expected_digest,
            ),
        )
    authority, active_path = authority_match
    facts = replace(
        facts,
        active_template_origin=authority.origin,
        active_template_version=authority.version,
    )
    if not active_path.is_file():
        return _failure("missing", "source_not_found", "active template 文件不存在", facts, field="path")
    try:
        observed_digest = file_sha256(active_path)
    except OSError as exc:
        return _failure("missing", "source_unreadable", f"源文件不可读：{type(exc).__name__}", facts, field="path")
    facts = replace(facts, observed_digest=observed_digest)
    try:
        authority_digest = _normalize_sha256(
            authority.expected_active_digest,
            field="template_authorities.expected_active_digest",
            required=False,
        )
    except _RefSchemaError as exc:
        return _failure("invalid", exc.code, exc.detail, facts, field=exc.field)
    if authority_digest is not None and authority_digest != observed_digest:
        return _failure(
            "stale",
            "authority_digest_mismatch",
            "active template 与冻结 authority digest 不一致",
            facts,
            field="template_authorities.expected_active_digest",
        )
    if expected_digest != observed_digest:
        return _failure("stale", "digest_mismatch", "source_ref digest 与当前文件不一致", facts, field="digest")

    if kind == "xlsx":
        return _validate_xlsx(raw, active_path, facts)
    return _validate_docx(raw, active_path, facts)


def validate_source_ref(raw: Mapping[str, Any], context: SourceRefContext) -> SourceRefValidation:
    """验证单条来源引用；不抛业务失败，统一返回结构化状态。

    未知 kind / 生命周期名作 kind → fail-closed（经 G-ID registry）。
    """
    # 延迟导入避免模块加载环；registry 对 xlsx/docx 回调 ``_validate_template_bound_ref``。
    from app.services.guidance_gid import validate_source_ref_via_registry

    return validate_source_ref_via_registry(raw, context)


def validate_source_refs(raw: Any, context: SourceRefContext) -> SourceRefBatchValidation:
    """验证引用数组并生成稳定 facts digest；空数组显式返回 missing。"""
    if not isinstance(raw, Sequence) or isinstance(raw, (str, bytes, bytearray)):
        results = (
            _failure(
                "invalid",
                "ref_list_invalid",
                "source_refs 必须是数组",
                SourceRefFacts(),
                field="source_refs",
            ),
        )
    elif not raw:
        results = (
            _failure(
                "missing",
                "ref_list_empty",
                "source_refs 为空",
                SourceRefFacts(),
                field="source_refs",
            ),
        )
    else:
        results = tuple(validate_source_ref(item, context) for item in raw)
    counts = Counter(item.status for item in results)
    facts_payload = [
        {
            "status": item.status,
            "issues": [asdict(issue) for issue in item.issues],
            "facts": asdict(item.facts),
        }
        for item in results
    ]
    return SourceRefBatchValidation(
        results=results,
        status_counts=dict(sorted(counts.items())),
        facts_digest=_stable_digest(facts_payload),
    )


def __getattr__(name: str) -> Any:
    """模板 lineage 冻结已迁至 ``guidance_template_authority``；保留既有导入路径。

    延迟解析而非顶层 import —— 该模块反过来依赖本模块的 ``TemplateAuthority`` /
    ``SourceRefContext``，直接 import 会成环。
    """
    if name in {
        "TemplateAuthoritySnapshot",
        "build_template_authority_snapshot",
        "build_source_ref_contexts",
    }:
        from app.services import guidance_template_authority

        return getattr(guidance_template_authority, name)
    raise AttributeError(f"module {__name__!r} has no attribute {name!r}")

"""附注正文（富文本 HTML）→ Word 段落 / 表格：附注正文进 Word 的唯一渲染（design §5.5，第 18 组）。

``disclosure_notes.text_content`` 存 HTML（单章编辑器存 ``innerHTML``，批量写入写消毒后的 HTML），
而 Word 导出曾把它当一段纯文字写入 ⇒ ``<p>…</p><p>&nbsp;</p>`` 原样出现在 Word 里（第 17 组真实写入验收）。

四个出口共用本模块：程序化导出、模板占位填充、交付件增量刷新、旧版全套填充。
**纯文字按编辑器口径渲染**（第 22 组）：先经 :func:`note_text_to_html`（markdown，与编辑器加载、批量写入同一转换）
再按块输出 —— 空行分段、单换行是段内换行、markdown 表格成 Word 表格、编号列表续编号。首版纯文字「一段一 run」：
生成链 / 大模型写入的纯文字（真库 3 个项目各 100+ 章）在 Word 里并成一段，表格以竖线原样出现，与编辑器所见不一致。

映射（与编辑器所见一致，不引入新样式）：
- ``p`` / ``div`` / ``blockquote`` → 段落；空段（含追加分隔 ``<p>&nbsp;</p>``、``<p><br></p>``）→ 无 run 的空段；
- ``h1``~``h6`` → 加粗段落 —— **不用 Word 标题样式**：导出带自动目录，正文里的小标题不应进目录；
- ``ol`` → 「（1）」前缀（附注惯用；交付件回填把「1. 中文」开头的段当章节标题剔除，用「1. 」会误判已修改），
  ``ul`` → 「• 」，嵌套每层缩进两格；
- ``strong`` / ``b`` 加粗、``em`` / ``i`` 倾斜，``br`` 段内换行，实体解码，其余标签只保留文字；
- ``table`` → Word 表格（单线边框，``th`` 加粗，``colspan`` / ``rowspan`` 合并，跨度上限 :data:`MAX_SPAN`）。

:func:`note_text_plain` 给交付件回填 / 刷新比对用：与交付 docx 提取（``section_anchor_utils.block_text_of``：
逐段 ``w:t`` 拼接、跳过表格与空段、段内换行不保留）同口径，守卫逐样本对拍「渲染后再提取」。
"""

from __future__ import annotations

import copy
import re
from collections.abc import Callable
from dataclasses import dataclass, field
from html.parser import HTMLParser
from typing import Any

from docx.oxml import OxmlElement
from docx.oxml.ns import qn
from docx.text.paragraph import Paragraph

from app.services.note_html_sanitizer import is_html_note_text, note_text_to_html

__all__ = [
    "MAX_SPAN",
    "append_note_text",
    "fill_paragraph_with_note_text",
    "is_html_note_text",
    "note_text_elements",
    "note_text_plain",
]

#: 单元格跨行 / 跨列上限：消毒器放行 1~999，原样建 999 列的表会拖垮导出
MAX_SPAN = 50

# 「正文是 HTML」的判据（is_html_note_text）在 note_html_sanitizer：批量写入与 Word 渲染共用一个口径
_BLOCK_TAGS = frozenset({"p", "div", "blockquote", "li", "h1", "h2", "h3", "h4", "h5", "h6"})
_HEADING_TAGS = frozenset({"h1", "h2", "h3", "h4", "h5", "h6"})
_DROP_TAGS = frozenset({"script", "style", "iframe", "object", "embed", "template", "svg", "math", "head", "title"})
_WS = re.compile(r"[ \t\r\n\f]+")

RunStyle = Callable[[Any, bool], None]
ParagraphStyle = Callable[[Any], None]


@dataclass
class _Run:
    text: str
    bold: bool = False
    italic: bool = False


@dataclass
class _Para:
    runs: list[_Run] = field(default_factory=list)
    heading: bool = False
    prefix: str = ""

    def has_text(self) -> bool:
        return any(r.text for r in self.runs)


@dataclass
class _Cell:
    paragraphs: list[_Para] = field(default_factory=list)
    header: bool = False
    colspan: int = 1
    rowspan: int = 1


@dataclass
class _Table:
    rows: list[list[_Cell]] = field(default_factory=list)


def _span(attrs: list[tuple[str, str | None]], name: str) -> int:
    raw = dict(attrs).get(name) or "1"
    return max(1, min(int(raw), MAX_SPAN)) if raw.isdigit() else 1


def _list_start(attrs: list[tuple[str, str | None]]) -> int:
    """``<ol start="N">`` 的起始编号（续编号列表）；缺省 / 非法按 1。"""
    raw = (dict(attrs).get("start") or "1").strip()
    return int(raw) if raw.isdigit() and 0 < int(raw) < 10000 else 1


def _as_html(text: str) -> str:
    """渲染入口：HTML 原样解析；纯文字先按编辑器口径转 HTML（markdown，第 22 组）。"""
    return text if is_html_note_text(text) else note_text_to_html(text)


class _BlockParser(HTMLParser):
    """HTML → 块列表（段落 / 表格）。容错：未闭合标签在 :meth:`close` 时收尾。"""

    def __init__(self) -> None:
        super().__init__(convert_charrefs=True)
        self.blocks: list[_Para | _Table] = []
        self._para: _Para | None = None
        self._para_sink: list | None = None
        self._bold = self._italic = self._drop = 0
        self._lists: list[list[Any]] = []  # [kind, counter]
        self._tables: list[dict[str, Any]] = []  # {"table": _Table, "row": list|None, "cell": _Cell|None}

    # ── 段落 ──
    def _sink(self) -> list:
        ctx = self._tables[-1] if self._tables else None
        if ctx is not None:
            if ctx["cell"] is None:  # 表格内、单元格外的文字：补一个单元格承接，不丢字
                self._open_cell(ctx, header=False, colspan=1, rowspan=1)
            return ctx["cell"].paragraphs
        return self.blocks

    def _open_para(self, *, heading: bool = False, prefix: str = "") -> _Para:
        self._para, self._para_sink = _Para(heading=heading, prefix=prefix), self._sink()
        return self._para

    def _flush(self) -> None:
        if self._para is not None and self._para_sink is not None:
            _trim(self._para)
            self._para_sink.append(self._para)
        self._para = self._para_sink = None

    def _list_prefix(self) -> str:
        kind_counter = self._lists[-1]
        kind_counter[1] += 1
        indent = "  " * (len(self._lists) - 1)
        return f"{indent}（{kind_counter[1]}）" if kind_counter[0] == "ol" else f"{indent}• "

    # ── 表格 ──
    def _open_cell(self, ctx: dict[str, Any], *, header: bool, colspan: int, rowspan: int) -> None:
        self._close_cell(ctx)
        if ctx["row"] is None:
            ctx["row"] = []
        ctx["cell"] = _Cell(header=header, colspan=colspan, rowspan=rowspan)

    def _close_cell(self, ctx: dict[str, Any]) -> None:
        if ctx["cell"] is not None:
            if self._para_sink is ctx["cell"].paragraphs:
                self._flush()
            ctx["row"].append(ctx["cell"])
            ctx["cell"] = None

    def _close_row(self, ctx: dict[str, Any]) -> None:
        self._close_cell(ctx)
        if ctx["row"]:
            ctx["table"].rows.append(ctx["row"])
        ctx["row"] = None

    def _close_table(self) -> None:
        ctx = self._tables.pop()
        self._close_row(ctx)
        if not ctx["table"].rows:
            return
        sink = self._sink()
        if sink is self.blocks:
            sink.append(ctx["table"])
        else:  # 嵌套表格：Word 单元格里不再嵌表，逐行拍平成段落
            for row in ctx["table"].rows:
                runs = [r for cell in row for p in cell.paragraphs for r in (p.runs + [_Run(" ")])]
                sink.append(_trim(_Para(runs=runs[:-1])))

    # ── HTMLParser 回调 ──
    def handle_starttag(self, tag: str, attrs: list[tuple[str, str | None]]) -> None:
        if tag in _DROP_TAGS:
            self._drop += 1
            return
        if self._drop:
            return
        if tag in ("strong", "b"):
            self._bold += 1
        elif tag in ("em", "i"):
            self._italic += 1
        elif tag == "br":
            (self._para or self._open_para()).runs.append(_Run("\n", self._bold > 0, self._italic > 0))
        elif tag in ("ul", "ol"):
            self._flush()
            self._lists.append([tag, _list_start(attrs) - 1 if tag == "ol" else 0])
        elif tag in _BLOCK_TAGS:
            prefix = self._list_prefix() if tag == "li" and self._lists else ""
            if self._para is not None and not self._para.has_text() and not prefix:
                self._para.heading = self._para.heading or tag in _HEADING_TAGS  # 嵌套块（如 li>p）沿用同一段
                return
            self._flush()
            self._open_para(heading=tag in _HEADING_TAGS, prefix=prefix)
        elif tag == "hr":
            self._flush()
        elif tag == "table":
            self._flush()
            self._tables.append({"table": _Table(), "row": None, "cell": None})
        elif tag == "tr" and self._tables:
            self._close_row(self._tables[-1])
            self._tables[-1]["row"] = []
        elif tag in ("td", "th") and self._tables:
            self._open_cell(self._tables[-1], header=tag == "th",
                            colspan=_span(attrs, "colspan"), rowspan=_span(attrs, "rowspan"))

    def handle_startendtag(self, tag: str, attrs: list[tuple[str, str | None]]) -> None:
        self.handle_starttag(tag, attrs)
        if tag not in ("br", "hr"):
            self.handle_endtag(tag)

    def handle_endtag(self, tag: str) -> None:
        if tag in _DROP_TAGS:
            self._drop = max(0, self._drop - 1)
            return
        if self._drop:
            return
        if tag in ("strong", "b"):
            self._bold = max(0, self._bold - 1)
        elif tag in ("em", "i"):
            self._italic = max(0, self._italic - 1)
        elif tag in ("ul", "ol"):
            self._flush()
            if self._lists:
                self._lists.pop()
        elif tag in _BLOCK_TAGS:
            self._flush()
        elif tag in ("td", "th") and self._tables:
            self._close_cell(self._tables[-1])
        elif tag == "tr" and self._tables:
            self._close_row(self._tables[-1])
        elif tag == "table" and self._tables:
            self._flush()
            self._close_table()

    def handle_data(self, data: str) -> None:
        if self._drop:
            return
        text = _WS.sub(" ", data)
        if self._para is None:
            if not text.strip(" "):
                return  # 块与块之间的排版空白
            self._open_para()
        para = self._para
        if para.runs and para.runs[-1].text.endswith("\n"):
            text = text.lstrip(" ")  # 换行后的行首空白浏览器不显示（如 <br />\n 后的排版换行）
            if not text:
                return
        bold, italic = self._bold > 0 or para.heading, self._italic > 0
        if para.runs and (para.runs[-1].bold, para.runs[-1].italic) == (bold, italic):
            para.runs[-1].text += text
        else:
            para.runs.append(_Run(text, bold, italic))

    def close(self) -> None:
        super().close()
        self._flush()
        while self._tables:
            self._close_table()


def _trim(para: _Para) -> _Para:
    """段首段尾空白去掉（段首换行保留：那是用户可见的空行）；全空白（含 &nbsp;）⇒ 空段。"""
    if para.runs:
        para.runs[0].text = para.runs[0].text.lstrip(" ")
        para.runs[-1].text = para.runs[-1].text.rstrip(" \n")
    para.runs = [r for r in para.runs if r.text]
    if not "".join(r.text for r in para.runs).strip():
        para.runs = []
    return para


def _parse(html: str) -> list[_Para | _Table]:
    parser = _BlockParser()
    parser.feed(html)
    parser.close()
    return parser.blocks


# ── Word 输出 ─────────────────────────────────────────────────────────────────


def _write_runs(paragraph: Paragraph, para: _Para, run_style: RunStyle | None, *, force_bold: bool = False) -> None:
    if not para.runs:
        return
    if para.prefix:
        para = _Para(runs=[_Run(para.prefix, para.runs[0].bold, para.runs[0].italic), *para.runs])
    for item in para.runs:
        run = paragraph.add_run(item.text)
        bold = item.bold or force_bold
        if run_style is not None:
            run_style(run, bold)
        elif bold:
            run.bold = True
        if item.italic:
            run.italic = True


def _set_borders(table: Any) -> None:
    borders = OxmlElement("w:tblBorders")
    for edge in ("top", "left", "bottom", "right", "insideH", "insideV"):
        el = OxmlElement(f"w:{edge}")
        for key, value in (("val", "single"), ("sz", "4"), ("space", "0"), ("color", "000000")):
            el.set(qn(f"w:{key}"), value)
        borders.append(el)
    table._tbl.tblPr.append(borders)


def _add_table(document: Any, spec: _Table, run_style: RunStyle | None) -> Any:
    """按占位网格摆放单元格（rowspan 占住下方格子），先合并再填字；返回 python-docx Table（在文末）。"""
    occupied: set[tuple[int, int]] = set()
    placements: list[tuple[int, int, _Cell]] = []
    for r, row in enumerate(spec.rows):
        c = 0
        for cell in row:
            while (r, c) in occupied:
                c += 1
            cell.rowspan = min(cell.rowspan, len(spec.rows) - r)  # 跨出表尾的行不存在（与浏览器一致）
            occupied.update((r + dr, c + dc) for dr in range(cell.rowspan) for dc in range(cell.colspan))
            placements.append((r, c, cell))
            c += cell.colspan
    n_rows = max(r for r, _ in occupied) + 1
    n_cols = max(c for _, c in occupied) + 1
    table = document.add_table(rows=n_rows, cols=n_cols)
    _set_borders(table)
    targets = []
    for r, c, cell in placements:
        target = table.cell(r, c)
        if cell.rowspan > 1 or cell.colspan > 1:
            target = target.merge(table.cell(r + cell.rowspan - 1, c + cell.colspan - 1))
        targets.append((target, cell))
    for target, cell in targets:
        for i, para in enumerate(cell.paragraphs):
            paragraph = target.paragraphs[0] if i == 0 else target.add_paragraph()
            _write_runs(paragraph, para, run_style, force_bold=cell.header)
    return table


BlockFilter = Callable[[list[Any]], list[Any]]


def _blocks(text: str | None, block_filter: BlockFilter | None) -> list[_Para | _Table]:
    """解析正文；``block_filter``（交付过滤，``note_deliverable_filter.make_block_filter``）作用于解析结果。"""
    blocks = _parse(_as_html(text or ""))
    return block_filter(blocks) if block_filter is not None else blocks


def append_note_text(document: Any, text: str | None, *, run_style: RunStyle | None = None,
                     paragraph_style: ParagraphStyle | None = None,
                     block_filter: BlockFilter | None = None) -> None:
    """文末追加正文（程序化导出 / 旧版填充 / 交付件刷新）：HTML 与纯文字同一渲染（纯文字先转 HTML）；全空白不追加。"""
    if not text or not text.strip():
        return
    for block in _blocks(text, block_filter):
        if isinstance(block, _Table):
            _add_table(document, block, run_style)
            continue
        paragraph = document.add_paragraph()
        _write_runs(paragraph, block, run_style)
        if paragraph_style is not None:
            paragraph_style(paragraph)


def fill_paragraph_with_note_text(paragraph: Paragraph, text: str, *, run_style: RunStyle | None = None,
                                  block_filter: BlockFilter | None = None) -> None:
    """模板占位段原位填充（占位符独占一段时用）：首块写进占位段，其余块（含表格）依次插在其后，
    新段落沿用占位段的段落属性；首块是表格时占位段留作空段。"""
    paragraph.clear()
    ppr = paragraph._p.pPr
    last = paragraph._p
    for index, block in enumerate(_blocks(text, block_filter)):
        if isinstance(block, _Table):
            element = _add_table(paragraph.part.document, block, run_style)._tbl  # 先建在文末，再移到位
        elif index == 0:
            _write_runs(paragraph, block, run_style)
            continue
        else:
            element = OxmlElement("w:p")
            if ppr is not None:
                element.append(copy.deepcopy(ppr))
            _write_runs(Paragraph(element, paragraph._parent), block, run_style)
        last.addnext(element)
        last = element


def note_text_elements(text: str, *, run_style: RunStyle | None = None,
                       block_filter: BlockFilter | None = None) -> list[Any]:
    """渲染成脱离文档的 ``w:p`` / ``w:tbl`` 元素（交付件刷新插入用）；不带样式引用，可移入任意文档。"""
    import docx

    scratch = docx.Document()
    body = scratch.element.body
    content_tags = (qn("w:p"), qn("w:tbl"))
    # 按位置取新增元素：新内容都追加在既有内容之后（sectPr 之前）。🔴 不能按 id() 比对 ——
    # lxml 的元素代理对象每次访问都可能新建，id 不稳定（首版因此丢了首段）
    before = sum(1 for el in body if el.tag in content_tags)
    append_note_text(scratch, text, run_style=run_style, block_filter=block_filter)
    elements = [el for el in body if el.tag in content_tags][before:]
    for el in elements:
        body.remove(el)
    return elements


def note_text_plain(text: str | None) -> str:
    """比对用文字：纯文字原样返回；HTML 按交付 docx 提取口径投影（逐段、跳过表格与空段、段内换行不保留）。

    回填比对两侧共用：交付 docx 提取出的文字是纯文字，不能再按 markdown 解释一遍 ⇒ 纯文字原样。
    要「库里纯文字正文渲染后的文字」用 ``note_text_plain(note_text_to_html(text))``。
    """
    if not is_html_note_text(text):
        return text or ""
    lines: list[str] = []
    for block in _parse(text or ""):
        if isinstance(block, _Table) or not block.runs:
            continue
        line = (block.prefix + "".join(r.text for r in block.runs)).replace("\n", "").strip()
        if line and not line.startswith("##"):
            lines.append(line)
    return "\n".join(lines)

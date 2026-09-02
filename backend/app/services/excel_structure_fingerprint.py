"""Excel OOXML 结构指纹 / identity inventory / visible-equivalence 报告（可复用模块）。

spec: workpaper-html-onlyoffice-bidirectional-writeback-closure
Requirements 6.13 / 6.14 / 6.15 / 6.16 / 6.17 / 6.20
Task 5（真实 OO 9.4 Excel identity 黑盒 probe）产出，Task 17（instrumentation definition
与 non-current upgrade candidate 生成器）与 Task 36/37 直接 import 复用。

为什么必须是模块而不是一次性脚本：
  Requirement 6.17 的「instrumentation 前后可见 sheet/业务值/公式/样式/merge/drawing/chart/pivot
  等价」在 Task 17 的 upgrader 里是**发布前置校验**，在 Task 5 里是 probe 判据；两处必须用
  同一套判据，否则 probe 通过的载体在 upgrader 上可能用不同口径重新判一次（假绿第③源）。

三层判据分工：
  1. `structure_fingerprint()` —— 只描述**事实**（sheet/defined name/table/merge/formula/
     style/part digest），不做任何通过与否的判断。
  2. `identity_inventory()` —— 从事实里投影出三类 identity 载体的清册（Property 66 的比对对象）。
  3. `visible_equivalence_report()` —— 两份 workbook 的逐 aspect 等价判定（Requirement 6.17）。

🔴 禁止 fail-open：任何采集失败都抛 `FingerprintError`，或在 `errors` 里留 ERROR 记录并让
调用方守卫打红。绝不把解析失败降级成「无数据 / 空 inventory」（那会让「载体被剥离」与
「采集器坏了」两种完全不同的事实产生同一个绿色结论）。
"""

from __future__ import annotations

import hashlib
import io
import re
import zipfile
from dataclasses import dataclass, field
from typing import Any
from xml.etree import ElementTree as ET

__all__ = [
    "FingerprintError",
    "GT_SYNC_SHEET_NAME",
    "IDENTITY_CARRIERS",
    "WorkbookFingerprint",
    "fingerprint_bytes",
    "identity_inventory",
    "structure_fingerprint",
    "visible_equivalence_report",
    "EQUIVALENCE_ASPECTS",
]

#: 隐藏 metadata sheet 名（design §Excel instrumentation contract
#: `ignored_by_business_sheet_enumerators`）。业务 sheet 枚举必须显式排除。
GT_SYNC_SHEET_NAME = "_GT_SYNC"

#: design §Identity probe 决策门的三类 Excel 候选载体（顺序即 contract 声明顺序）。
IDENTITY_CARRIERS = ("hidden_sheet", "defined_name", "hidden_uuid_column")

#: Requirement 6.17 要求逐项比对的 aspect（visible-equivalence 报告的固定列）。
EQUIVALENCE_ASPECTS = (
    "visible_sheets",
    "business_values",
    "formulas",
    "styles",
    "merges",
    "protected_parts",
)

_NS = {
    "main": "http://schemas.openxmlformats.org/spreadsheetml/2006/main",
    "r": "http://schemas.openxmlformats.org/officeDocument/2006/relationships",
    "ct": "http://schemas.openxmlformats.org/package/2006/content-types",
    "pr": "http://schemas.openxmlformats.org/package/2006/relationships",
}

#: 「受保护部件」= 平台不管理、instrumentation 与 OO 往返都不得改变的 OOXML 部件类别。
#: Requirement 6.17 的 drawing/chart/pivot 就落在这里。
_PROTECTED_PART_PATTERNS: dict[str, re.Pattern[str]] = {
    "drawing": re.compile(r"^xl/drawings/"),
    "chart": re.compile(r"^xl/charts/"),
    "pivot_table": re.compile(r"^xl/pivotTables/"),
    "pivot_cache": re.compile(r"^xl/pivotCache/"),
    "media": re.compile(r"^xl/media/"),
    "vba": re.compile(r"^xl/vbaProject\.bin$"),
    "external_link": re.compile(r"^xl/externalLinks/"),
    "slicer": re.compile(r"^xl/slicers?/"),
    "query_table": re.compile(r"^xl/queryTables/"),
    "conn": re.compile(r"^xl/connections\.xml$"),
}


class FingerprintError(RuntimeError):
    """结构采集失败。**不得**被调用方降级为「无数据」。"""


def _sha256(data: bytes) -> str:
    return hashlib.sha256(data).hexdigest()


def _hash_lines(lines: list[str]) -> str:
    return _sha256("\n".join(lines).encode("utf-8"))


def _local(tag: str) -> str:
    return tag.rsplit("}", 1)[-1]


def _col_letters(ref: str) -> str:
    match = re.match(r"^([A-Z]+)", ref)
    if not match:
        raise FingerprintError(f"非法单元格引用: {ref!r}")
    return match.group(1)


def _col_index(letters: str) -> int:
    """A→1, Z→26, AA→27。"""
    idx = 0
    for ch in letters:
        idx = idx * 26 + (ord(ch) - 64)
    return idx


def _row_number(ref: str) -> int:
    match = re.search(r"(\d+)$", ref)
    if not match:
        raise FingerprintError(f"非法单元格引用: {ref!r}")
    return int(match.group(1))


# ---------------------------------------------------------------------------
# 数据结构
# ---------------------------------------------------------------------------


@dataclass
class WorkbookFingerprint:
    """一份 xlsx 的结构事实。字段全部是**描述**，不含通过/失败判断。"""

    byte_sha256: str
    byte_len: int
    #: [{name, sheet_id, state, rel_target}]，顺序即 workbook.xml 中的 tab 顺序
    sheets: list[dict[str, Any]]
    #: [{name, scope, ref, hidden}]；scope=None 为 workbook 级，否则为 sheet 名
    defined_names: list[dict[str, Any]]
    #: [{sheet, part, name, display_name, ref, columns:[...], header_row}]
    tables: list[dict[str, Any]]
    #: {sheet: [merge ranges]}
    merges: dict[str, list[str]]
    #: {sheet: {coord: formula}}
    formulas: dict[str, dict[str, str]]
    #: {sheet: {coord: 值的 repr}}（只含非空）
    cell_values: dict[str, dict[str, str]]
    #: {sheet: {coord: 样式指纹}}
    cell_styles: dict[str, dict[str, str]]
    #: {sheet: [被隐藏的列字母]}
    hidden_columns: dict[str, list[str]]
    #: {sheet: [被隐藏的行号]}
    hidden_rows: dict[str, list[int]]
    #: {part_name: sha256}，全 zip 条目
    part_digests: dict[str, str]
    #: {类别: [part_name]}，见 `_PROTECTED_PART_PATTERNS`
    protected_parts: dict[str, list[str]]
    #: 采集过程中的非致命异常（**必须为空**才算采集成功；守卫据此打红）
    errors: list[str] = field(default_factory=list)

    # -- 派生投影（供守卫/报告直接比对，避免各处自己再算一遍） -------------

    @property
    def sheet_names(self) -> list[str]:
        return [s["name"] for s in self.sheets]

    @property
    def visible_sheet_names(self) -> list[str]:
        return [s["name"] for s in self.sheets if s["state"] == "visible"]

    @property
    def hidden_sheet_names(self) -> list[str]:
        return [s["name"] for s in self.sheets if s["state"] != "visible"]

    def business_sheet_names(self, *, ignored: tuple[str, ...] = (GT_SYNC_SHEET_NAME,)) -> list[str]:
        """业务 sheet 枚举口径：可见且不在 metadata 排除名单内（Requirement 6.17 后半句）。"""
        return [name for name in self.visible_sheet_names if name not in ignored]

    def aspect_digests(self, *, ignored_sheets: tuple[str, ...] = (GT_SYNC_SHEET_NAME,)) -> dict[str, str]:
        """Requirement 6.17 的六个 aspect 各自的 digest（只覆盖业务 sheet）。"""
        business = self.business_sheet_names(ignored=ignored_sheets)
        return {
            "visible_sheets": _hash_lines(business),
            "business_values": _hash_lines(
                [
                    f"{sheet}!{coord}={value}"
                    for sheet in business
                    for coord, value in sorted(self.cell_values.get(sheet, {}).items())
                ]
            ),
            "formulas": _hash_lines(
                [
                    f"{sheet}!{coord}={formula}"
                    for sheet in business
                    for coord, formula in sorted(self.formulas.get(sheet, {}).items())
                ]
            ),
            "styles": _hash_lines(
                [
                    f"{sheet}!{coord}={style}"
                    for sheet in business
                    for coord, style in sorted(self.cell_styles.get(sheet, {}).items())
                ]
            ),
            "merges": _hash_lines(
                [f"{sheet}!{rng}" for sheet in business for rng in sorted(self.merges.get(sheet, []))]
            ),
            "protected_parts": _hash_lines(
                [
                    f"{category}:{part}={self.part_digests.get(part, 'MISSING')}"
                    for category, parts in sorted(self.protected_parts.items())
                    for part in sorted(parts)
                ]
            ),
        }

    def to_json(self) -> dict[str, Any]:
        """可入 evidence 的投影（大字典折成 digest + 计数，保留可复核的结构清册）。"""
        return {
            "byte_sha256": self.byte_sha256,
            "byte_len": self.byte_len,
            "sheets": self.sheets,
            "visible_sheet_names": self.visible_sheet_names,
            "hidden_sheet_names": self.hidden_sheet_names,
            "defined_names": self.defined_names,
            "tables": self.tables,
            "merges": {k: v for k, v in sorted(self.merges.items()) if v},
            "hidden_columns": {k: v for k, v in sorted(self.hidden_columns.items()) if v},
            "hidden_rows": {k: v for k, v in sorted(self.hidden_rows.items()) if v},
            "formula_counts": {k: len(v) for k, v in sorted(self.formulas.items()) if v},
            "nonempty_cell_counts": {k: len(v) for k, v in sorted(self.cell_values.items()) if v},
            "protected_parts": {k: v for k, v in sorted(self.protected_parts.items()) if v},
            "part_count": len(self.part_digests),
            "part_digests": self.part_digests,
            "aspect_digests": self.aspect_digests(),
            "errors": self.errors,
        }


# ---------------------------------------------------------------------------
# 采集：zip / workbook.xml 层
# ---------------------------------------------------------------------------


def _parse_workbook_xml(zf: zipfile.ZipFile) -> tuple[list[dict[str, Any]], list[dict[str, Any]]]:
    """从 `xl/workbook.xml` + rels 读 sheet 清册与 defined names。

    直接读 XML 而不用 openpyxl 的原因：openpyxl 会丢弃 `veryHidden` 之外的一些属性，
    且 `wb.defined_names` 在 3.1 里对 sheet-local 名称的暴露方式几经变化；probe 要的是
    **磁盘上的原始事实**。
    """
    try:
        wb_xml = zf.read("xl/workbook.xml")
    except KeyError as exc:
        raise FingerprintError("缺少 xl/workbook.xml，不是合法 xlsx") from exc
    try:
        rels_xml = zf.read("xl/_rels/workbook.xml.rels")
    except KeyError as exc:
        raise FingerprintError("缺少 xl/_rels/workbook.xml.rels") from exc

    rel_map: dict[str, str] = {}
    for node in ET.fromstring(rels_xml):
        if _local(node.tag) != "Relationship":
            continue
        rel_map[node.attrib["Id"]] = node.attrib.get("Target", "")

    root = ET.fromstring(wb_xml)
    sheets: list[dict[str, Any]] = []
    for node in root.findall("main:sheets/main:sheet", _NS):
        rid = node.attrib.get(f"{{{_NS['r']}}}id", "")
        target = rel_map.get(rid, "")
        sheets.append(
            {
                "name": node.attrib.get("name", ""),
                "sheet_id": node.attrib.get("sheetId", ""),
                "state": node.attrib.get("state", "visible"),
                "rel_id": rid,
                "rel_target": target,
            }
        )

    defined_names: list[dict[str, Any]] = []
    sheet_order = [s["name"] for s in sheets]
    for node in root.findall("main:definedNames/main:definedName", _NS):
        local_id = node.attrib.get("localSheetId")
        scope = None
        if local_id is not None:
            try:
                scope = sheet_order[int(local_id)]
            except (ValueError, IndexError):
                scope = f"<invalid localSheetId {local_id}>"
        defined_names.append(
            {
                "name": node.attrib.get("name", ""),
                "scope": scope,
                "ref": (node.text or "").strip(),
                "hidden": node.attrib.get("hidden") in ("1", "true"),
            }
        )
    return sheets, defined_names


def _normalise_part(target: str) -> str:
    """rels Target → zip 条目路径。"""
    target = target.lstrip("/")
    if target.startswith("xl/"):
        return target
    return f"xl/{target}"


def _parse_tables(zf: zipfile.ZipFile, sheets: list[dict[str, Any]]) -> list[dict[str, Any]]:
    """逐 sheet 读 `<tableParts>` → `xl/tables/tableN.xml`，取 Table identity 与列清册。"""
    tables: list[dict[str, Any]] = []
    for sheet in sheets:
        sheet_part = _normalise_part(sheet["rel_target"])
        if sheet_part not in zf.namelist():
            continue
        try:
            sheet_root = ET.fromstring(zf.read(sheet_part))
        except ET.ParseError as exc:
            raise FingerprintError(f"{sheet_part} 解析失败: {exc}") from exc
        parts_node = sheet_root.find("main:tableParts", _NS)
        if parts_node is None:
            continue
        rels_path = f"{sheet_part.rsplit('/', 1)[0]}/_rels/{sheet_part.rsplit('/', 1)[1]}.rels"
        rel_map: dict[str, str] = {}
        if rels_path in zf.namelist():
            for node in ET.fromstring(zf.read(rels_path)):
                if _local(node.tag) == "Relationship":
                    rel_map[node.attrib["Id"]] = node.attrib.get("Target", "")
        for part_ref in parts_node.findall("main:tablePart", _NS):
            rid = part_ref.attrib.get(f"{{{_NS['r']}}}id", "")
            target = rel_map.get(rid, "")
            table_part = target.replace("../", "xl/").lstrip("/")
            if not table_part.startswith("xl/"):
                table_part = _normalise_part(target)
            if table_part not in zf.namelist():
                raise FingerprintError(
                    f"{sheet['name']} 的 tablePart {rid} 指向不存在的部件 {table_part!r}"
                )
            table_root = ET.fromstring(zf.read(table_part))
            columns = [
                col.attrib.get("name", "")
                for col in table_root.findall("main:tableColumns/main:tableColumn", _NS)
            ]
            ref = table_root.attrib.get("ref", "")
            header_row = None
            if ref and table_root.attrib.get("headerRowCount", "1") != "0":
                header_row = _row_number(ref.split(":")[0])
            tables.append(
                {
                    "sheet": sheet["name"],
                    "part": table_part,
                    "id": table_root.attrib.get("id", ""),
                    "name": table_root.attrib.get("name", ""),
                    "display_name": table_root.attrib.get("displayName", ""),
                    "ref": ref,
                    "header_row": header_row,
                    "columns": columns,
                }
            )
    return tables


def _normalise_cell_value(value: Any) -> str:
    """把单元格值规范成**可比**字符串。

    🔴 不能直接 `repr()`：openpyxl 把数组公式包成 `ArrayFormula`、把数据表公式包成
    `DataTableFormula`，两者都没实现 `__repr__`/`__eq__` ⇒ `repr()` 里带内存地址
    （`<... object at 0x0000...>`），同一份文件读两次都不相等，等价报告会恒判「不等价」。
    实测：C24 的「参考-本福特定律测试」有 100 个数组公式，未规范化时 instrumentation
    前后凭空出现 100 处 diff。
    """
    from openpyxl.worksheet.formula import ArrayFormula, DataTableFormula

    if isinstance(value, ArrayFormula):
        return f"ArrayFormula(ref={value.ref!r},text={value.text!r})"
    if isinstance(value, DataTableFormula):
        return "DataTableFormula(" + ",".join(
            f"{attr}={getattr(value, attr, None)!r}"
            for attr in ("ref", "ca", "dt2D", "dtr", "r1", "r2", "del1", "del2")
        ) + ")"
    return repr(value)


def _is_formula(cell: Any, value: Any) -> bool:
    """公式判定：`data_type == 'f'`（含数组/数据表公式）或字符串以 `=` 开头。"""
    from openpyxl.worksheet.formula import ArrayFormula, DataTableFormula

    if isinstance(value, (ArrayFormula, DataTableFormula)):
        return True
    if getattr(cell, "data_type", None) == "f":
        return True
    return isinstance(value, str) and value.startswith("=")


def _color_sig(color: Any) -> str:
    """openpyxl Color → 稳定的 `(type:value)` 指纹。

    🔴 不能直接读 `.rgb`：openpyxl 的 `Color.rgb` 只在 `type=='rgb'` 时有值，
    对 indexed/theme/auto 色返回字符串 `"Values must be of type <class 'str'>"`
    （它的 descriptor 校验错误消息）。直接拼进指纹会造成两类假差异：
      - 同一份文件里 indexed 与 rgb 色被判成"不同"（其实一个是未设置）；
      - Task 5 实测：K11 只改一个格，styles 却报 444/767 处 diff，全是这个占位串。
    """
    if color is None:
        return "none"
    ctype = getattr(color, "type", None)
    if ctype == "rgb":
        return f"rgb:{color.rgb}"
    if ctype == "indexed":
        return f"indexed:{getattr(color, 'indexed', None)}"
    if ctype == "theme":
        return f"theme:{getattr(color, 'theme', None)}/{getattr(color, 'tint', 0)}"
    if ctype == "auto":
        return "auto"
    return f"{ctype}:unknown"


def _style_signature(cell: Any) -> str:
    """单元格样式指纹：数字格式 + 字体 + 填充 + 边框 + 对齐。

    只取**语义**属性（不含 style index），因为 OO 重排样式表时 index 会变但语义不变。
    """
    font = cell.font
    fill = cell.fill
    align = cell.alignment
    border = cell.border
    return "|".join(
        [
            f"nf={cell.number_format}",
            f"font={font.name},{font.sz},{int(bool(font.b))},{int(bool(font.i))},{_color_sig(font.color)}",
            f"fill={fill.patternType},{_color_sig(getattr(fill, 'fgColor', None))},"
            f"{_color_sig(getattr(fill, 'bgColor', None))}",
            f"align={align.horizontal},{align.vertical},{int(bool(align.wrapText))}",
            "border="
            + ",".join(
                str(getattr(getattr(border, side), "style", None))
                for side in ("left", "right", "top", "bottom")
            ),
        ]
    )


def structure_fingerprint(data: bytes) -> WorkbookFingerprint:
    """采集一份 xlsx 的全部结构事实。

    失败一律抛 `FingerprintError`（禁止返回空结构）。局部可容忍问题记入 `errors`，
    调用方守卫必须断言 `errors == []`。
    """
    if not data[:2] == b"PK":
        raise FingerprintError("不是 zip/xlsx 字节流（缺少 PK 魔数）")

    errors: list[str] = []
    with zipfile.ZipFile(io.BytesIO(data)) as zf:
        bad = zf.testzip()
        if bad is not None:
            raise FingerprintError(f"zip 条目损坏: {bad}")
        part_digests = {name: _sha256(zf.read(name)) for name in sorted(zf.namelist())}
        sheets, defined_names = _parse_workbook_xml(zf)
        tables = _parse_tables(zf, sheets)
        protected: dict[str, list[str]] = {}
        for category, pattern in _PROTECTED_PART_PATTERNS.items():
            hits = [name for name in sorted(part_digests) if pattern.match(name)]
            if hits:
                protected[category] = hits

    # 值 / 公式 / 样式 / merge / 隐藏行列走 openpyxl（只读不写，纯读取无副作用）。
    import openpyxl

    merges: dict[str, list[str]] = {}
    formulas: dict[str, dict[str, str]] = {}
    values: dict[str, dict[str, str]] = {}
    styles: dict[str, dict[str, str]] = {}
    hidden_cols: dict[str, list[str]] = {}
    hidden_rows: dict[str, list[int]] = {}
    try:
        wb = openpyxl.load_workbook(io.BytesIO(data), data_only=False, keep_vba=False)
    except Exception as exc:  # noqa: BLE001 — 转成显式 FingerprintError，绝不吞成空结构
        raise FingerprintError(f"openpyxl 无法打开: {type(exc).__name__}: {exc}") from exc
    try:
        for ws in wb.worksheets:
            title = ws.title
            merges[title] = [str(rng) for rng in ws.merged_cells.ranges]
            sheet_formulas: dict[str, str] = {}
            sheet_values: dict[str, str] = {}
            sheet_styles: dict[str, str] = {}
            for row in ws.iter_rows():
                for cell in row:
                    if cell.value is None:
                        continue
                    normalised = _normalise_cell_value(cell.value)
                    if _is_formula(cell, cell.value):
                        sheet_formulas[cell.coordinate] = normalised
                    else:
                        sheet_values[cell.coordinate] = normalised
                    try:
                        sheet_styles[cell.coordinate] = _style_signature(cell)
                    except Exception as exc:  # noqa: BLE001 — 记 ERROR，不静默跳过
                        errors.append(f"style {title}!{cell.coordinate}: {type(exc).__name__}: {exc}")
            formulas[title] = sheet_formulas
            values[title] = sheet_values
            styles[title] = sheet_styles
            hidden_cols[title] = sorted(
                (letter for letter, dim in ws.column_dimensions.items() if dim.hidden),
                key=_col_index,
            )
            hidden_rows[title] = sorted(idx for idx, dim in ws.row_dimensions.items() if dim.hidden)
    finally:
        wb.close()

    return WorkbookFingerprint(
        byte_sha256=_sha256(data),
        byte_len=len(data),
        sheets=sheets,
        defined_names=defined_names,
        tables=tables,
        merges=merges,
        formulas=formulas,
        cell_values=values,
        cell_styles=styles,
        hidden_columns=hidden_cols,
        hidden_rows=hidden_rows,
        part_digests=part_digests,
        protected_parts=protected,
        errors=errors,
    )


def fingerprint_bytes(data: bytes) -> dict[str, Any]:
    """`structure_fingerprint(...).to_json()` 的便捷入口。"""
    return structure_fingerprint(data).to_json()


# ---------------------------------------------------------------------------
# identity inventory（Property 66 的比对对象）
# ---------------------------------------------------------------------------


def _read_gt_sync_pairs(data: bytes, sheet_name: str = GT_SYNC_SHEET_NAME) -> dict[str, str]:
    """读隐藏 metadata sheet 的 key→value（A 列 key、B 列 value）。"""
    import openpyxl

    wb = openpyxl.load_workbook(io.BytesIO(data), data_only=False)
    try:
        if sheet_name not in wb.sheetnames:
            return {}
        ws = wb[sheet_name]
        pairs: dict[str, str] = {}
        for row in ws.iter_rows(min_col=1, max_col=2):
            key = row[0].value
            if key is None or str(key).strip() == "":
                continue
            value = row[1].value if len(row) > 1 else None
            pairs[str(key)] = "" if value is None else str(value)
        return pairs
    finally:
        wb.close()


def identity_inventory(
    data: bytes,
    *,
    expected_table: str | None = None,
    uuid_column_letter: str | None = None,
    uuid_column_header: str | None = None,
    uuid_sheet_id: str | None = None,
    uuid_sheet_name: str | None = None,
    metadata_sheet: str = GT_SYNC_SHEET_NAME,
    defined_name_prefix: str = "GT_",
) -> dict[str, Any]:
    """三类 identity 载体的实测清册。

    返回结构对每个载体给出 `present` + 具体内容 + `errors`。**不做通过判定** —— 判定由
    Task 5 契约 / Task 17 upgrader 按各自期望值比对（避免把期望写进采集器造成锁死错值）。

    UUID 列定位口径（Requirement 6.14/6.20）：
      1. 首选 `uuid_sheet_id` + `uuid_column_letter` —— sheetId 不随展示名变化，是**唯一**
         能在用户改 sheet 名后仍定位的锚点；
      2. 次选 `uuid_sheet_name` + `uuid_column_letter`（仅审计线索口径，用于对照实证展示名
         是否已漂移）；
      3. `uuid_column_header` 只在 Table `headerRowCount>0` 时可用（表头写在单元格里）。
    列内值**全列扫描**而不是按固定行范围取：插删行后 UUID 会整体位移，固定范围会把
    「位移」误判成「丢失」。
    """
    fp = structure_fingerprint(data)
    inv: dict[str, Any] = {
        "byte_sha256": fp.byte_sha256,
        "sheet_names": fp.sheet_names,
        "errors": list(fp.errors),
    }

    # --- 载体 1：hidden metadata sheet ---
    meta_sheet = next((s for s in fp.sheets if s["name"] == metadata_sheet), None)
    pairs = _read_gt_sync_pairs(data, metadata_sheet) if meta_sheet else {}
    inv["hidden_sheet"] = {
        "present": meta_sheet is not None,
        "sheet_name": metadata_sheet,
        "state": meta_sheet["state"] if meta_sheet else None,
        "is_hidden": bool(meta_sheet and meta_sheet["state"] in ("hidden", "veryHidden")),
        "pairs": pairs,
        "pair_count": len(pairs),
        "excluded_from_business_enumeration": metadata_sheet not in fp.business_sheet_names(),
    }

    # --- 载体 2：defined names ---
    gt_names = [d for d in fp.defined_names if d["name"].startswith(defined_name_prefix)]
    inv["defined_name"] = {
        "present": bool(gt_names),
        "prefix": defined_name_prefix,
        "names": {d["name"]: d["ref"] for d in gt_names},
        "scopes": {d["name"]: d["scope"] for d in gt_names},
        "hidden_flags": {d["name"]: d["hidden"] for d in gt_names},
        "count": len(gt_names),
        "all_defined_name_count": len(fp.defined_names),
    }

    # --- 载体 3a：Excel Table identity（与 UUID 列**分开**判定）---
    carrier_errors: list[str] = []
    table = None
    if expected_table:
        table = next(
            (t for t in fp.tables if expected_table in (t["name"], t["display_name"])),
            None,
        )
        if table is None and fp.tables:
            carrier_errors.append(
                f"期望 Table {expected_table!r} 不在实测清册中: "
                f"{[t['display_name'] for t in fp.tables]}"
            )
    inv["excel_table"] = {
        "expected_table": expected_table,
        "present": table is not None,
        "table_name": table["display_name"] if table else None,
        "table_ref": table["ref"] if table else None,
        "table_sheet": table["sheet"] if table else None,
        "table_part": table["part"] if table else None,
        "header_row_count_zero": (table["header_row"] is None) if table else None,
        "table_columns": table["columns"] if table else [],
        "all_tables": [
            {"display_name": t["display_name"], "sheet": t["sheet"], "ref": t["ref"]}
            for t in fp.tables
        ],
        "all_table_count": len(fp.tables),
    }

    # --- 载体 3b：隐藏 row UUID 列 ---
    # 🔴 三条 sheet 定位路径**全部**独立求解并各自报告，绝不静默 fallback：
    #    静默 fallback 会把「某个锚点已被 OO 破坏」掩盖成「定位成功」（假绿第①源）。
    #    Task 5 实证：OO 9.4 每次保存都把 sheetId 按 tab 顺序重编号 ⇒ sheet_id 路径会解析
    #    到**另一张 sheet**，此时 fallback 到 table_sheet 就会让守卫看不见这个事实。
    candidates: dict[str, dict[str, Any]] = {}
    if uuid_sheet_id is not None:
        match = next((s for s in fp.sheets if s["sheet_id"] == str(uuid_sheet_id)), None)
        candidates["sheet_id"] = {
            "requested": str(uuid_sheet_id),
            "resolved_sheet": match["name"] if match else None,
            "resolvable": match is not None,
        }
    if uuid_sheet_name is not None:
        candidates["sheet_name"] = {
            "requested": uuid_sheet_name,
            "resolved_sheet": uuid_sheet_name if uuid_sheet_name in fp.sheet_names else None,
            "resolvable": uuid_sheet_name in fp.sheet_names,
        }
    if table is not None:
        candidates["table_sheet"] = {
            "requested": table["display_name"],
            "resolved_sheet": table["sheet"],
            "resolvable": True,
        }

    col_letter = uuid_column_letter
    if col_letter is None and table is not None and uuid_column_header:
        if uuid_column_header not in table["columns"]:
            carrier_errors.append(
                f"table {table['display_name']!r} 的列清册不含 UUID 列 "
                f"{uuid_column_header!r}: {table['columns']}"
            )
        else:
            offset = table["columns"].index(uuid_column_header)
            start_col = _col_index(_col_letters(table["ref"].split(":")[0]))
            col_letter = _index_to_letters(start_col + offset)

    def _scan(sheet: str | None) -> tuple[dict[int, str], bool | None]:
        if not sheet or not col_letter:
            return {}, None
        found = {
            _row_number(coord): raw.strip("'\"")
            for coord, raw in fp.cell_values.get(sheet, {}).items()
            if _col_letters(coord) == col_letter
        }
        return found, col_letter in fp.hidden_columns.get(sheet, [])

    for name, info in candidates.items():
        found, hidden = _scan(info["resolved_sheet"])
        info["row_uuid_count"] = len(found)
        info["uuid_column_hidden"] = hidden

    # 「实际拿到 UUID 的那条路径」用于报告数据本体是否存活；判定各锚点是否可靠仍看 candidates。
    winner = next(
        (
            name
            for name in ("sheet_id", "sheet_name", "table_sheet")
            if candidates.get(name, {}).get("row_uuid_count")
        ),
        None,
    )
    sheet_name = candidates[winner]["resolved_sheet"] if winner else (
        candidates.get("sheet_id", {}).get("resolved_sheet")
        or candidates.get("sheet_name", {}).get("resolved_sheet")
        or candidates.get("table_sheet", {}).get("resolved_sheet")
    )
    row_uuids, col_hidden = _scan(sheet_name)
    values = list(row_uuids.values())
    declared_anchor_ok = bool(candidates.get("sheet_id", {}).get("row_uuid_count"))
    if uuid_sheet_id is not None and not candidates["sheet_id"]["resolvable"]:
        carrier_errors.append(
            f"sheetId {uuid_sheet_id!r} 在实测 workbook 中不存在: "
            f"{[(s['name'], s['sheet_id']) for s in fp.sheets]}"
        )
    inv["hidden_uuid_column"] = {
        "present": bool(row_uuids),
        "resolved_sheet": sheet_name,
        "resolved_sheet_by": winner,
        "sheet_resolution_candidates": candidates,
        #: sheetId 路径能否直接定位到 UUID 列（Requirement 6.14 的候选锚点是否成立）
        "sheet_id_anchor_holds": declared_anchor_ok,
        "requested_sheet_id": uuid_sheet_id,
        "requested_sheet_name": uuid_sheet_name,
        "sheet_display_name_drifted": bool(
            uuid_sheet_name and sheet_name and uuid_sheet_name != sheet_name
        ),
        "uuid_column_letter": col_letter,
        "uuid_column_header": uuid_column_header,
        "uuid_column_hidden": col_hidden,
        "row_uuids": {str(k): v for k, v in sorted(row_uuids.items())},
        "row_uuid_count": len(row_uuids),
        "distinct_row_uuid_count": len(set(values)),
        "empty_row_uuids": sorted((str(k) for k, v in row_uuids.items() if not v), key=int),
        "duplicate_row_uuids": sorted({v for v in values if values.count(v) > 1 and v}),
        "errors": carrier_errors,
    }
    inv["errors"].extend(carrier_errors)
    return inv


def _index_to_letters(index: int) -> str:
    """1→A, 27→AA。"""
    letters = ""
    while index > 0:
        index, rem = divmod(index - 1, 26)
        letters = chr(65 + rem) + letters
    return letters


# ---------------------------------------------------------------------------
# visible-equivalence（Requirement 6.17）
# ---------------------------------------------------------------------------


def visible_equivalence_report(
    before: bytes,
    after: bytes,
    *,
    ignored_sheets: tuple[str, ...] = (GT_SYNC_SHEET_NAME,),
    allow_added_hidden_columns: dict[str, list[str]] | None = None,
    allow_added_table_columns: dict[str, list[str]] | None = None,
    label_before: str = "before",
    label_after: str = "after",
) -> dict[str, Any]:
    """instrumentation（或 OO 往返）前后的可见等价报告。

    Requirement 6.17：可见 sheet、业务值、公式、样式、merge、drawing/chart/pivot 等
    受保护部件必须等价；隐藏 metadata sheet 由 `ignored_sheets` 显式排除。

    `allow_added_hidden_columns` / `allow_added_table_columns` 是**受控白名单**：
    instrumentation 会在业务 sheet 上新增一个隐藏 UUID 列（Requirement 6.13 明确允许），
    该列的新增单元格不算破坏等价；除白名单外的任何差异都判不等价。
    """
    fp_before = structure_fingerprint(before)
    fp_after = structure_fingerprint(after)
    allow_cols = {k: set(v) for k, v in (allow_added_hidden_columns or {}).items()}
    allow_table_cols = {k: set(v) for k, v in (allow_added_table_columns or {}).items()}

    business_before = fp_before.business_sheet_names(ignored=ignored_sheets)
    business_after = fp_after.business_sheet_names(ignored=ignored_sheets)

    def _filtered(mapping: dict[str, dict[str, str]], sheet: str, *, allowed: set[str]) -> dict[str, str]:
        return {
            coord: val
            for coord, val in mapping.get(sheet, {}).items()
            if _col_letters(coord) not in allowed
        }

    aspects: dict[str, Any] = {}

    # 1. visible sheets（含顺序）
    aspects["visible_sheets"] = {
        "equivalent": business_before == business_after,
        label_before: business_before,
        label_after: business_after,
        "added": [s for s in business_after if s not in business_before],
        "removed": [s for s in business_before if s not in business_after],
    }

    # 2/3/4. 业务值、公式、样式（逐业务 sheet 比对，排除受控新增列）
    for aspect, source in (
        ("business_values", "cell_values"),
        ("formulas", "formulas"),
        ("styles", "cell_styles"),
    ):
        diffs: list[dict[str, Any]] = []
        for sheet in business_before:
            allowed = allow_cols.get(sheet, set())
            b = _filtered(getattr(fp_before, source), sheet, allowed=allowed)
            a = _filtered(getattr(fp_after, source), sheet, allowed=allowed)
            for coord in sorted(set(b) | set(a)):
                if b.get(coord) != a.get(coord):
                    diffs.append(
                        {
                            "sheet": sheet,
                            "coord": coord,
                            label_before: b.get(coord),
                            label_after: a.get(coord),
                        }
                    )
        aspects[aspect] = {
            "equivalent": not diffs,
            "diff_count": len(diffs),
            "first_diffs": diffs[:20],
        }

    # 5. merge
    merge_diffs: list[dict[str, Any]] = []
    for sheet in business_before:
        b = sorted(fp_before.merges.get(sheet, []))
        a = sorted(fp_after.merges.get(sheet, []))
        if b != a:
            merge_diffs.append(
                {
                    "sheet": sheet,
                    "only_" + label_before: [r for r in b if r not in a],
                    "only_" + label_after: [r for r in a if r not in b],
                }
            )
    aspects["merges"] = {
        "equivalent": not merge_diffs,
        "diff_count": len(merge_diffs),
        "first_diffs": merge_diffs[:20],
    }

    # 6. 受保护部件（drawing/chart/pivot/media/vba/...）：部件名与字节 digest 双比
    part_diffs: list[dict[str, Any]] = []
    categories = sorted(set(fp_before.protected_parts) | set(fp_after.protected_parts))
    for category in categories:
        b_parts = fp_before.protected_parts.get(category, [])
        a_parts = fp_after.protected_parts.get(category, [])
        if sorted(b_parts) != sorted(a_parts):
            part_diffs.append({"category": category, "only_before": [p for p in b_parts if p not in a_parts],
                               "only_after": [p for p in a_parts if p not in b_parts]})
            continue
        for part in sorted(b_parts):
            if fp_before.part_digests.get(part) != fp_after.part_digests.get(part):
                part_diffs.append(
                    {
                        "category": category,
                        "part": part,
                        label_before: fp_before.part_digests.get(part),
                        label_after: fp_after.part_digests.get(part),
                    }
                )
    aspects["protected_parts"] = {
        "equivalent": not part_diffs,
        "categories_present": {c: len(fp_before.protected_parts.get(c, [])) for c in categories},
        "diff_count": len(part_diffs),
        "first_diffs": part_diffs[:20],
        "covered": bool(categories),
    }

    # 表结构新增列白名单核对（instrumentation 只允许加声明过的列）
    table_col_diffs: list[dict[str, Any]] = []
    before_tables = {(t["sheet"], t["display_name"]): t for t in fp_before.tables}
    after_tables = {(t["sheet"], t["display_name"]): t for t in fp_after.tables}
    for key, t_after in after_tables.items():
        t_before = before_tables.get(key)
        if t_before is None:
            continue
        added = [c for c in t_after["columns"] if c not in t_before["columns"]]
        unexpected = [c for c in added if c not in allow_table_cols.get(key[0], set())]
        if unexpected:
            table_col_diffs.append({"sheet": key[0], "table": key[1], "unexpected_added_columns": unexpected})

    hidden_sheets_added = [s for s in fp_after.hidden_sheet_names if s not in fp_before.hidden_sheet_names]
    report = {
        "label_before": label_before,
        "label_after": label_after,
        "ignored_sheets": list(ignored_sheets),
        "byte_sha256": {label_before: fp_before.byte_sha256, label_after: fp_after.byte_sha256},
        "aspects": aspects,
        "aspect_digests": {
            label_before: fp_before.aspect_digests(ignored_sheets=ignored_sheets),
            label_after: fp_after.aspect_digests(ignored_sheets=ignored_sheets),
        },
        "hidden_sheets_added": hidden_sheets_added,
        "metadata_sheet_excluded_from_business": all(
            name not in business_after for name in ignored_sheets
        ),
        "unexpected_table_columns": table_col_diffs,
        "collection_errors": {label_before: fp_before.errors, label_after: fp_after.errors},
    }
    report["equivalent"] = (
        all(aspects[a]["equivalent"] for a in EQUIVALENCE_ASPECTS)
        and not table_col_diffs
        and not fp_before.errors
        and not fp_after.errors
    )
    report["aspect_verdicts"] = {a: bool(aspects[a]["equivalent"]) for a in EQUIVALENCE_ASPECTS}
    return report

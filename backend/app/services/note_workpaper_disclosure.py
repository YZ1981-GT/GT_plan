"""底稿披露 sheet 切表：按推送注册表找底稿册里的附注披露 sheet，切出逐张表（表名候选、逐列表头段、表头下的行名）。
只读，不连库。

spec: consol-note-header-library-and-node-reaggregate R8.4（design §17.5）；单体附注模板 ↔ 底稿的配对与逐列比对见
``note_template_workpaper_alignment``。口径（用户 2026-10-23）：

- 【…】是给用户的提示：从每格文字里去掉（未闭合的【到格尾；A 列上下几格连着写的一段提示整段去掉，见
  ``_drop_multi_cell_hints``）；行名 / 组合名按实际情况填、只作参考，不比。
- 表头行 = B 列起有非公式、非占位文字的连续行；表名 = 表头上方最近的「只有 A 列有字」的行（表名候选取到上一张表的数据行）。
- 表头格按合并区域展开；相邻网格列在全部表头行属于同一格 ⇒ 一个逻辑列。一格只盖一个逻辑列、上下几行拆开写的
  （「利息资 / 本化累 / 计金额」）连成一段；盖住多个逻辑列的才是分组。
- 列路径里去掉：勾稽注记（「（1）」「C=A+B」）、括号提示（「（不适用删除）」）；表头为空或末段是「……」的列不算
  （空白间隔列 / 按需增加的占位列）。
- 表头格是单格引用公式（「='被投资单位财务信息（合营、联营）G7-5'!C8」，引用格是文字）的，取被引用格的字；标签列
  的合并格从更上面一行开始的（A229:B230，第 229 行其余格只有引用公式），表头从合并格首行算起 —— G7 国企合营 / 联营
  财务信息表的企业名行原先因此丢掉（R8.4 批 9b-2）。只认表头行里的引用；数据行不解析。
"""
from __future__ import annotations

import json
import re
import unicodedata
from dataclasses import dataclass, field
from functools import lru_cache
from pathlib import Path
from typing import Any

from app.services.note_table_header_library import norm

BACKEND = Path(__file__).resolve().parents[2]
DATA_DIR = BACKEND / "data"
REGISTRY_PATH = DATA_DIR / "note_workpaper_sync_registry.json"
LIBRARY_PATH = DATA_DIR / "gt_template_library.json"
STDS = ("listed", "soe")

#: 推送注册表（由前端常量生成）里底稿 sheet 名缺失 / 与册里不符的：(底稿, 口径) → 册里实际 sheet 名。
#: 注册表的 sheet 名在册里找得到时不用这里（失效即报，见 ``stale_sheet_overrides``）。
SHEET_OVERRIDES: dict[tuple[str, str], str] = {
    ("L1", "soe"): "附注披露信息核对（国企）",
    ("L3", "soe"): "附注披露（国企）信息核对",
    ("M10", "listed"): "附注披露信息核对（上市公司）",
    ("M10", "soe"): "附注披露信息核对（国企）",
}

_HINT = re.compile(r"【[^】]*】|【.*$", re.S)
_WS = re.compile(r"[\s\u3000]+")
_PLACEHOLDER = re.compile(r"^(?:[A-Z](?:=.*)?|[(（]?X+(?:\.X+)?[)）]?|/|—+|-+|…+|\.+|0)$")
_BANNER = ("索引号", "索引：", "页次", "页  次", "致同会计师事务所")
_NOTE_SEG = re.compile(r"^(?:[（(]\d{1,2}[)）](?:[=＝].*)?|[A-Z](?:[=＝].*)?)$")
_HINT_SEG = re.compile(r"^[（(][^）)]*(?:不适用|删除|可删|根据|注[:：]|选择)[^）)]*[）)]$")
_TAIL_NOTE = re.compile(r"[（(]\d{1,2}[)）]$")
_ELLIPSIS = frozenset({"……", "…", "...", "......", "………"})
_NAME_HINT = re.compile(r"[（(][^（）()]*(?:不适用|删除|注[:：]|如无|可删|根据|选择|格式)[^（）()]*[）)]")
_LEAD = re.compile(r"^(?:[（(]\s*[0-9一二三四五六七八九十]{1,3}\s*[）)]|[0-9]{1,2}[、.．]|[一二三四五六七八九十]{1,3}[、.．]|[①-⑳])")
_TAIL = re.compile(r"[：:。]+$")
_MAX_ROW_LABELS = 40


def cell_text(value: Any) -> str:
    """单元格显示文字：公式 / 非文本为空；去【】提示与首尾空白，去换行；「返回目录」链接格当空格。"""
    if not isinstance(value, str):
        return ""
    s = value.strip()
    if s.startswith("="):
        return ""
    s = _HINT.sub("", s).replace("\n", "").replace("\r", "").strip()
    return "" if _WS.sub("", s) == "返回目录" else s


def _nk(text: Any) -> str:
    return _WS.sub("", unicodedata.normalize("NFKC", str(text or "")))


def name_key(text: Any) -> str:
    """表名比较键：去【】与括号提示、行首序号、尾部冒号。行首序号在 NFKC 前后各去一次：NFKC 把「①」变成「1」，
    之后就认不出是序号了（「① 按账龄披露」曾得到「1按账龄披露」）。"""
    t = _WS.sub("", _NAME_HINT.sub("", _HINT.sub("", str(text or ""))))
    for _ in range(2):
        t = _LEAD.sub("", t)
    t = _nk(t)
    for _ in range(2):
        t = _LEAD.sub("", t)
    return _TAIL.sub("", t)


def is_continuation(name: Any) -> bool:
    k = name_key(name)
    return k.startswith("续") or "(续" in k


def base_name(name: Any) -> str:
    k = re.sub(r"\(续[^)]*\)$", "", name_key(name))
    return re.sub(r"\((期末余额|上年年末余额|期末数|期初数|期末|期初)\)$", "", k)


# ─── 底稿披露 sheet 切表 ─────────────────────────────────────────────────────


@dataclass(frozen=True)
class WpColumn:
    raw: tuple[str, ...]  # 自上而下各段原字（已去提示）
    cover: tuple[tuple[int, int], ...]  # 各段所在单元格盖住的逻辑列（首、末序号，含标签列）

    @property
    def path(self) -> tuple[str, ...]:
        return tuple(norm(x) for x in self.raw)


@dataclass
class WpTable:
    wp_code: str
    sheet: str
    header_rows: list[int]
    title: str
    title_candidates: list[str]
    columns: list[WpColumn]
    dropped: list[str] = field(default_factory=list)
    row_labels: list[str] = field(default_factory=list)  # 表头下面各行 A 列文字（行名只作参考：配对用，不比）

    @property
    def paths(self) -> list[tuple[str, ...]]:
        return [c.path for c in self.columns]


_SHEET_REF = re.compile(r"^=\s*(?:'((?:[^']|'')+)'|([^'!=()\s]+))!\$?([A-Z]{1,3})\$?(\d+)\s*$")
_LOCAL_REF = re.compile(r"^=\s*\$?([A-Z]{1,3})\$?(\d+)\s*$")


def ref_text(ws: Any, value: Any) -> str:
    """单格引用公式（跨 sheet 或本 sheet）→ 被引用格的显示文字（``cell_text`` 口径）；不是单格引用、引用格不是文字
    （数值 / 又一层公式 / 空）的返回空串。"""
    if not (isinstance(value, str) and value.lstrip().startswith("=")):
        return ""
    m = _SHEET_REF.match(value.strip())
    if m:
        title = (m.group(1) or "").replace("''", "'") or m.group(2)
        col, row = m.group(3), m.group(4)
    else:
        m = _LOCAL_REF.match(value.strip())
        if not m:
            return ""
        title, col, row = ws.title, m.group(1), m.group(2)
    book = getattr(ws, "parent", None)
    if book is None or title not in book.sheetnames:
        return ""
    return cell_text(book[title][f"{col}{row}"].value)


def _resolve_header_refs(ws: Any, vals: dict, owner: dict, rows: list[int], maxc: int) -> dict:
    """``vals`` 副本：表头行里单格引用公式的格换成被引用格的字（只换表头行，数据行不动）。"""
    out = dict(vals)
    for r in rows:
        for c in range(1, maxc + 1):
            o = owner.get((r, c), (r, c))
            if o not in out or cell_text(out[o]):
                continue
            text = ref_text(ws, out[o])
            if text:
                out[o] = text
    return out


def _grid(ws: Any) -> tuple[dict, dict]:
    vals = {(c.row, c.column): c.value for row in ws.iter_rows() for c in row if c.value not in (None, "")}
    owner: dict[tuple[int, int], tuple[int, int]] = {}
    for rng in ws.merged_cells.ranges:
        for r in range(rng.min_row, rng.max_row + 1):
            for c in range(rng.min_col, rng.max_col + 1):
                owner[(r, c)] = (rng.min_row, rng.min_col)
    return vals, owner


def _cell(vals: dict, owner: dict, r: int, c: int) -> tuple[tuple[int, int], str]:
    o = owner.get((r, c), (r, c))
    return o, cell_text(vals.get(o))


def _header_like(vals: dict, owner: dict, r: int, maxc: int) -> bool:
    for c in range(2, maxc + 1):
        o, t = _cell(vals, owner, r, c)
        if o[1] != 1 and t and not _PLACEHOLDER.match(_WS.sub("", t)):
            return True
    return False


def _a_only(vals: dict, owner: dict, r: int, maxc: int) -> bool:
    if not cell_text(vals.get(owner.get((r, 1), (r, 1)))):
        return False
    return not any(_cell(vals, owner, r, c)[1] for c in range(2, maxc + 1) if _cell(vals, owner, r, c)[0][1] != 1)


def _blank(vals: dict, owner: dict, r: int, maxc: int) -> bool:
    return not any(_cell(vals, owner, r, c)[1] for c in range(1, maxc + 1))


def _has_data(vals: dict, r: int, maxc: int) -> bool:
    """B 列起有公式或数值 ⇒ 数据行（表名候选不越过它）。"""
    for c in range(2, maxc + 1):
        v = vals.get((r, c))
        if isinstance(v, (int, float)) or (isinstance(v, str) and v.strip().startswith("=")):
            return True
    return False


def _brackets(text: str) -> tuple[int, int]:
    """(没闭合的「【」数, 没有开头的「】」数)。"""
    depth = orphan = 0
    for ch in text:
        if ch == "【":
            depth += 1
        elif ch == "】":
            depth, orphan = (depth - 1, orphan) if depth else (0, orphan + 1)
    return depth, orphan


def _drop_multi_cell_hints(vals: dict, owner: dict, maxc: int) -> dict:
    """A 列上下几格连着写的一段提示（「【注：1.…；」「2.…」…「4.…。】」，每行只有 A 列有字、中间不空行）：开头一格
    ``cell_text`` 本就去掉「【」起的字，其后各格到收尾的「】」也是提示 —— 返回去掉这些字的 ``vals`` 副本（K6 国企首表
    原以提示末行「4.应披露…。】」为表名，配给了模板里同名的错表）。没有收尾、中间断开的不算一段，只去开头一格。"""
    out = dict(vals)

    def a_raw(r: int) -> str:
        v = vals.get(owner.get((r, 1), (r, 1)))
        return v if isinstance(v, str) else ""

    for r in sorted(r for (r, c) in vals if c == 1):
        if not _brackets(a_raw(r))[0]:
            continue
        k = r + 1
        while _a_only(vals, owner, k, maxc) and not _brackets(a_raw(k))[0]:
            if _brackets(a_raw(k))[1]:
                for x in range(r + 1, k):
                    out[owner.get((x, 1), (x, 1))] = ""
                out[owner.get((k, 1), (k, 1))] = a_raw(k)[a_raw(k).rindex("】") + 1:]
                break
            k += 1
    return out


def _extend_unmerged_groups(vals: dict, owner: dict, rows: list[int], last: int) -> dict[tuple[int, int], tuple[int, int]]:
    """表头格的归属：合并区域按左上角；**没合并、只写在首格的分组名**（底稿常见：「期末数」只写在 B 格、C / D 留空）
    盖住右侧同行的空格 —— 条件：不是最后一行表头、该空格不在任何合并区域里、它下面一行有字、左邻（数据列）是有字的
    单格（左邻是合并区域时以合并范围为准，不外延）。"""
    own: dict[tuple[int, int], tuple[int, int]] = {}
    merged_heads = set(owner.values())
    for i, hr in enumerate(rows):
        for c in range(1, last + 1):
            o, txt = _cell(vals, owner, hr, c)
            own[(hr, c)] = o
            if (txt or i == len(rows) - 1 or c <= 2 or (hr, c) in owner
                    or not _cell(vals, owner, rows[i + 1], c)[1]):
                continue
            left = own[(hr, c - 1)]
            if left[1] >= 2 and left not in merged_heads and cell_text(vals.get(left)):
                own[(hr, c)] = left
    return own


def extract_tables(ws: Any, *, wp_code: str = "", sheet: str = "") -> list[WpTable]:
    vals, owner = _grid(ws)
    maxr, maxc = ws.max_row, ws.max_column
    vals = _drop_multi_cell_hints(vals, owner, maxc)
    tables: list[WpTable] = []
    r, prev_end = 1, 0
    while r <= maxr:
        if not _header_like(vals, owner, r, maxc):
            r += 1
            continue
        h0 = r
        while r <= maxr and _header_like(vals, owner, r, maxc):
            r += 1
        top = owner.get((h0, 1), (h0, 1))[0]
        if prev_end < top < h0:
            h0 = top  # 标签列合并格从更上面一行开始：那几行也是表头（其余格只有引用公式，如企业名行）
        rows = list(range(h0, r))
        hv = _resolve_header_refs(ws, vals, owner, rows, maxc)
        if any(b in _cell(hv, owner, hr, c)[1] for hr in rows for c in range(1, maxc + 1) for b in _BANNER):
            prev_end = r - 1
            continue
        cands, t = [], h0 - 1
        while t > prev_end and len(cands) < 4 and not _has_data(vals, t, maxc):
            if _a_only(vals, owner, t, maxc):
                cands.append(cell_text(vals.get(owner.get((t, 1), (t, 1)))))
            elif not _blank(vals, owner, t, maxc):
                break
            t -= 1
        used = [c for c in range(1, maxc + 1) if any(_cell(hv, owner, hr, c)[1] for hr in rows)]
        last = max(used)
        own = _extend_unmerged_groups(hv, owner, rows, last)
        owners = [tuple(own[(hr, c)] for hr in rows) for c in range(1, last + 1)]
        logical = [i for i in range(last) if i == 0 or owners[i] != owners[i - 1]]
        # 每个单元格盖住几个逻辑列（分组判定）
        cover: dict[tuple[int, int], set[int]] = {}
        for li, gi in enumerate(logical):
            for hr_i in range(len(rows)):
                cover.setdefault(owners[gi][hr_i], set()).add(li)
        columns, dropped = [], []
        for li, gi in enumerate(logical):
            segs: list[tuple[str, tuple[int, int]]] = []
            seen: set[tuple[int, int]] = set()
            for hr_i, hr in enumerate(rows):
                o = owners[gi][hr_i]
                if o in seen:
                    continue
                seen.add(o)
                txt = _TAIL_NOTE.sub("", cell_text(hv.get(o))).strip()
                if txt and not _NOTE_SEG.match(_WS.sub("", txt)) and not _HINT_SEG.match(_WS.sub("", txt)):
                    span = cover.get(o, {li})
                    segs.append((txt, (min(span), max(span))))
            raw = tuple(x for x, _ in segs)
            if li > 0 and (not raw or _WS.sub("", raw[-1]) in _ELLIPSIS):
                dropped.append("/".join(raw) or "<空>")
                continue
            columns.append(WpColumn(raw=raw, cover=tuple(c for _, c in segs)))
        cands = [c for c in cands if c]
        labels_below: list[str] = []
        k = r
        while k <= maxr and len(labels_below) < _MAX_ROW_LABELS and not _header_like(vals, owner, k, maxc):
            if (txt := cell_text(vals.get(owner.get((k, 1), (k, 1))))):
                labels_below.append(txt)
            k += 1
        tables.append(WpTable(wp_code=wp_code, sheet=sheet, header_rows=rows, title=cands[0] if cands else "",
                              title_candidates=cands, columns=columns, dropped=dropped, row_labels=labels_below))
        prev_end = r - 1
    return tables


# ─── 找底稿 sheet ────────────────────────────────────────────────────────────


@lru_cache(maxsize=1)
def _registry() -> list[dict[str, Any]]:
    return json.loads(REGISTRY_PATH.read_text(encoding="utf-8"))["entries"]


@lru_cache(maxsize=1)
def _library() -> list[dict[str, Any]]:
    return json.loads(LIBRARY_PATH.read_text(encoding="utf-8"))["templates"]


def workbook_paths(wp_code: str) -> list[Path]:
    pat = re.compile(rf"^{re.escape(wp_code)}(?![0-9A-Za-z])")
    return [BACKEND / t["file_path"] for t in _library()
            if pat.match(t["code"]) and t["file_path"].lower().endswith((".xlsx", ".xlsm")) and (BACKEND / t["file_path"]).exists()]


@lru_cache(maxsize=64)
def _book(path: Path) -> Any:
    import openpyxl

    return openpyxl.load_workbook(path, data_only=False)


def find_sheet(wp_code: str, std: str, sheet: str | None) -> tuple[Path, str] | None:
    """注册表 sheet 名（去空白 / NFKC 后相等）→ 册里的 sheet；找不到再用 ``SHEET_OVERRIDES``。"""
    for name in (sheet, SHEET_OVERRIDES.get((wp_code, std))):
        if not name:
            continue
        for path in workbook_paths(wp_code):
            for sn in _book(path).sheetnames:
                if _nk(sn) == _nk(name):
                    return path, sn
    return None


def stale_sheet_overrides() -> list[str]:
    """登记了兜底却用不上（注册表的 sheet 名已能找到 / 兜底名在册里也没有）的条目。"""
    out = []
    reg = {(e["wp_code"], std): e.get(f"sheet_{std}") for e in _registry() for std in STDS}
    for (wp, std), name in SHEET_OVERRIDES.items():
        listed = reg.get((wp, std))
        if listed and any(_nk(sn) == _nk(listed) for p in workbook_paths(wp) for sn in _book(p).sheetnames):
            out.append(f"{wp}/{std}: 注册表的 sheet「{listed}」已找得到")
        elif not any(_nk(sn) == _nk(name) for p in workbook_paths(wp) for sn in _book(p).sheetnames):
            out.append(f"{wp}/{std}: 册里没有「{name}」")
    return out


def section_entries() -> dict[tuple[str, str], list[dict[str, Any]]]:
    """(口径, 附注章节号) → 推这一节的注册表条目（同一节几张底稿）。"""
    out: dict[tuple[str, str], list[dict[str, Any]]] = {}
    for e in _registry():
        for std in STDS:
            if e.get(std):
                out.setdefault((std, e[std]), []).append(e)
    return out


@lru_cache(maxsize=256)
def workpaper_tables(wp_code: str, std: str, sheet: str | None) -> tuple[WpTable, ...]:
    hit = find_sheet(wp_code, std, sheet)
    if hit is None:
        return ()
    path, sn = hit
    return tuple(extract_tables(_book(path)[sn], wp_code=wp_code, sheet=f"{path.name}!{sn}"))

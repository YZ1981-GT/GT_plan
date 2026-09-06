"""M5 盈余公积 — 业务逻辑服务

提供功能：
- 导出模板 / 导出数据 / 导入数据（xlsx三级）
- 法定盈余公积计提测试（10%计提验证 + 50%上限判断）

科目4101盈余公积（**贷方/权益类！**）：期末=期初+贷方-借方
法定盈余公积按净利润10%计提，累计达注册资本50%可不再计提。
任意盈余公积由股东大会决议计提。

Requirements: 6.1, 6.6
"""

from __future__ import annotations

import io
import json
import logging
from typing import Any
from uuid import uuid4

import sqlalchemy as sa
from openpyxl import Workbook, load_workbook
from openpyxl.styles import Alignment, Font, PatternFill
from sqlalchemy.ext.asyncio import AsyncSession

logger = logging.getLogger(__name__)

_M5_ACCOUNT_CODE = "4101"
_ROW_LIMIT = 500
_TOLERANCE = 0.01


# ═══════════════════════════════════════════════════════════════════════════════
# 导入导出配置
# ═══════════════════════════════════════════════════════════════════════════════

_SHEET_CONFIGS: dict[str, dict[str, Any]] = {
    "M5-2": {
        "title": "M5-2 明细表（法定+任意盈余公积）",
        "headers": [
            "类别", "项目名称", "期初余额", "本期计提(贷方)",
            "本期转增资本(借方)", "本期弥补亏损(借方)", "期末余额",
            "变动原因", "关联底稿", "备注",
        ],
        "fields": [
            "category", "item_name", "begin", "accrual",
            "capital_transfer", "loss_offset", "end_balance",
            "reason", "linked_wp", "remark",
        ],
    },
    "M5-4": {
        "title": "M5-4 计提检查表",
        "headers": [
            "检查项目", "净利润", "弥补以前年度亏损",
            "计提基数", "法定比例(%)", "应计提金额",
            "账面已计提", "差异", "备注",
        ],
        "fields": [
            "check_item", "net_profit", "prior_year_loss",
            "accrual_base", "rate_percent", "estimated_accrual",
            "booked_accrual", "diff", "remark",
        ],
    },
}

_HEADER_FONT = Font(bold=True, size=11)
_HEADER_FILL = PatternFill(start_color="4472C4", end_color="4472C4", fill_type="solid")
_HEADER_FONT_WHITE = Font(bold=True, size=11, color="FFFFFF")
_HEADER_ALIGN = Alignment(horizontal="center", vertical="center", wrap_text=True)


# ═══════════════════════════════════════════════════════════════════════════════
# 辅助函数
# ═══════════════════════════════════════════════════════════════════════════════


def _parse_num(v: Any) -> float:
    """安全解析数值，None/空/NaN→0.0."""
    if v is None or v == "":
        return 0.0
    try:
        f = float(v)
        return 0.0 if f != f else f
    except (ValueError, TypeError):
        return 0.0


async def _get_wp_context(wp_id: str, db: AsyncSession) -> dict[str, Any]:
    """从 wp_id 取关联的 project_id."""
    result = await db.execute(
        sa.text("""
            SELECT wp.project_id, p.audit_year
            FROM working_paper wp
            JOIN projects p ON p.id = wp.project_id
            WHERE wp.id = :wp_id
        """),
        {"wp_id": wp_id},
    )
    row = result.fetchone()
    if not row:
        raise ValueError(f"底稿不存在: {wp_id}")
    return {"project_id": row.project_id, "year": int(row.audit_year or 2025)}


def _create_sheet_with_headers(wb: Workbook, config: dict[str, Any]) -> None:
    """在工作簿中创建带格式列头的sheet."""
    ws = wb.create_sheet(title=config["title"])
    for col_idx, header in enumerate(config["headers"], 1):
        cell = ws.cell(row=1, column=col_idx, value=header)
        cell.font = _HEADER_FONT_WHITE
        cell.fill = _HEADER_FILL
        cell.alignment = _HEADER_ALIGN
    # 设置列宽
    for col_idx in range(1, len(config["headers"]) + 1):
        ws.column_dimensions[ws.cell(row=1, column=col_idx).column_letter].width = 16


# ═══════════════════════════════════════════════════════════════════════════════
# 导出模板
# ═══════════════════════════════════════════════════════════════════════════════


async def export_template(
    wp_id: str,
    db: AsyncSession,
    *,
    sheet: str | None = None,
) -> io.BytesIO:
    """导出空白xlsx模板（含列头格式）。

    Args:
        wp_id: 底稿ID
        db: 数据库会话
        sheet: 可选，指定导出单个sheet（M5-2/M5-4）
    """
    wb = Workbook()
    wb.remove(wb.active)  # type: ignore[arg-type]

    if sheet and sheet in _SHEET_CONFIGS:
        _create_sheet_with_headers(wb, _SHEET_CONFIGS[sheet])
    else:
        for cfg in _SHEET_CONFIGS.values():
            _create_sheet_with_headers(wb, cfg)

    buffer = io.BytesIO()
    wb.save(buffer)
    buffer.seek(0)
    return buffer


# ═══════════════════════════════════════════════════════════════════════════════
# 导出数据
# ═══════════════════════════════════════════════════════════════════════════════


async def export_data(
    wp_id: str,
    db: AsyncSession,
    *,
    sheet: str | None = None,
) -> io.BytesIO:
    """导出含当前数据的xlsx。

    从 checklist_responses 读取 M5 相关数据填入对应sheet。
    """
    wb = Workbook()
    wb.remove(wb.active)  # type: ignore[arg-type]

    # 读取 checklist_responses 数据
    responses = await _load_responses(wp_id, db)

    configs_to_export = (
        {sheet: _SHEET_CONFIGS[sheet]} if sheet and sheet in _SHEET_CONFIGS else _SHEET_CONFIGS
    )

    for sheet_key, cfg in configs_to_export.items():
        ws = wb.create_sheet(title=cfg["title"])
        # 写入列头
        for col_idx, header in enumerate(cfg["headers"], 1):
            cell = ws.cell(row=1, column=col_idx, value=header)
            cell.font = _HEADER_FONT_WHITE
            cell.fill = _HEADER_FILL
            cell.alignment = _HEADER_ALIGN

        # 写入数据行
        row_idx = 2
        sheet_suffix = sheet_key.split("-")[1] if "-" in sheet_key else sheet_key
        prefix = f"M5-{sheet_suffix}-row-"
        for item_id, data in sorted(responses.items()):
            if not item_id.startswith(prefix):
                continue
            raw = data.get("conclusion", "")
            if not raw:
                continue
            try:
                row_data = json.loads(raw) if isinstance(raw, str) else raw
            except (json.JSONDecodeError, TypeError):
                continue
            if not isinstance(row_data, dict):
                continue

            for col_idx, field in enumerate(cfg["fields"], 1):
                val = row_data.get(field, "")
                ws.cell(row=row_idx, column=col_idx, value=val)
            row_idx += 1

    buffer = io.BytesIO()
    wb.save(buffer)
    buffer.seek(0)
    return buffer


# ═══════════════════════════════════════════════════════════════════════════════
# 导入数据
# ═══════════════════════════════════════════════════════════════════════════════


async def import_data(
    wp_id: str,
    content: bytes,
    db: AsyncSession,
    *,
    sheet: str | None = None,
) -> dict[str, Any]:
    """解析xlsx并写入checklist_responses。

    返回:
        { imported_count: int, warnings: list[str] }
    """
    warnings: list[str] = []
    imported_count = 0

    try:
        wb = load_workbook(io.BytesIO(content), read_only=True, data_only=True)
    except Exception as e:
        raise ValueError(f"无法解析xlsx文件: {e}")

    # 确定要导入的sheet配置
    configs_to_import: dict[str, dict[str, Any]] = {}
    if sheet and sheet in _SHEET_CONFIGS:
        configs_to_import = {sheet: _SHEET_CONFIGS[sheet]}
    else:
        configs_to_import = _SHEET_CONFIGS.copy()

    for sheet_key, cfg in configs_to_import.items():
        # 在工作簿中查找匹配的sheet
        ws = None
        for ws_name in wb.sheetnames:
            if cfg["title"] in ws_name or sheet_key in ws_name:
                ws = wb[ws_name]
                break

        if ws is None:
            warnings.append(f"未找到sheet: {cfg['title']}")
            continue

        # 验证列头
        actual_headers = [
            str(ws.cell(row=1, column=c).value or "").strip()
            for c in range(1, len(cfg["headers"]) + 1)
        ]
        expected_headers = cfg["headers"]
        if actual_headers != expected_headers:
            match_count = sum(1 for a, e in zip(actual_headers, expected_headers) if a == e)
            if match_count < len(expected_headers) * 0.6:
                warnings.append(
                    f"Sheet {sheet_key} 列头不匹配: "
                    f"期望 {expected_headers[:3]}..., 实际 {actual_headers[:3]}..."
                )
                continue

        # 解析数据行
        sheet_suffix = sheet_key.split("-")[1] if "-" in sheet_key else sheet_key
        for row_idx in range(2, min(ws.max_row or 2, _ROW_LIMIT) + 1):
            row_data: dict[str, Any] = {}
            has_data = False
            for col_idx, field in enumerate(cfg["fields"], 1):
                val = ws.cell(row=row_idx, column=col_idx).value
                if val is not None and str(val).strip():
                    has_data = True
                row_data[field] = val

            if not has_data:
                continue

            item_id = f"M5-{sheet_suffix}-row-{row_idx - 1}"
            await db.execute(
                sa.text("""
                    INSERT INTO checklist_responses (id, wp_id, item_id, conclusion, status)
                    VALUES (:id, :wid, :iid, :conclusion, 'imported')
                    ON CONFLICT (wp_id, item_id)
                    DO UPDATE SET conclusion = :conclusion, status = 'imported'
                """),
                {
                    "id": str(uuid4()),
                    "wid": wp_id,
                    "iid": item_id,
                    "conclusion": json.dumps(row_data, ensure_ascii=False),
                },
            )
            imported_count += 1

    wb.close()
    return {"imported_count": imported_count, "warnings": warnings}


# ═══════════════════════════════════════════════════════════════════════════════
# 法定盈余公积计提测试
# ═══════════════════════════════════════════════════════════════════════════════


async def run_accrual_test(
    wp_id: str,
    request: Any,
    db: AsyncSession,
) -> Any:
    """法定盈余公积计提测试（10%计提验证 + 50%上限判断）。

    验证逻辑：
    1. 计提基数 = 净利润 - 弥补以前年度亏损
    2. 应计提 = 计提基数 × 法定比例(10%)
    3. 差异 = 应计提 - 账面已计提
    4. 累计法定盈余公积 ≥ 注册资本50% → 可不再计提
    """
    from app.routers.m5_surplus_reserve import AccrualTestResponse

    warnings: list[str] = []

    # 获取净利润（优先用请求参数，否则尝试从M6底稿读取）
    net_profit = request.net_profit
    if net_profit is None:
        net_profit = await _fetch_m6_net_profit(wp_id, db)
        if net_profit is None:
            net_profit = 0.0
            warnings.append("未获取到M6净利润数据，计提基数默认为0")

    # 核心计提计算
    prior_year_loss = request.prior_year_loss
    accrual_base = net_profit - prior_year_loss
    rate = request.rate
    estimated_accrual = accrual_base * rate if accrual_base > 0 else 0.0
    booked_accrual = request.booked_accrual
    diff = estimated_accrual - booked_accrual

    # 50%上限判断
    accumulated_reserve = request.accumulated_reserve
    registered_capital = request.registered_capital
    ceiling_threshold = registered_capital * 0.5
    ceiling_reached = accumulated_reserve >= ceiling_threshold if registered_capital > 0 else False

    # 合规性判断
    is_compliant = True
    if ceiling_reached:
        # 达到上限后可不再计提，任何计提额都合规
        if booked_accrual > 0:
            warnings.append("累计法定盈余公积已达注册资本50%，可不再计提，但本期仍有计提")
    else:
        # 未达上限，校验差异
        if abs(diff) > _TOLERANCE:
            is_compliant = False
            if diff > 0:
                warnings.append(f"计提不足：应计提{estimated_accrual:.2f}，实际{booked_accrual:.2f}，差异{diff:.2f}")
            else:
                warnings.append(f"超额计提：应计提{estimated_accrual:.2f}，实际{booked_accrual:.2f}，差异{diff:.2f}")

    if accrual_base < 0:
        warnings.append("计提基数为负（净利润<弥补亏损），无需计提法定盈余公积")

    return AccrualTestResponse(
        net_profit=net_profit,
        prior_year_loss=prior_year_loss,
        accrual_base=accrual_base,
        rate=rate,
        estimated_accrual=estimated_accrual,
        booked_accrual=booked_accrual,
        diff=round(diff, 2),
        accumulated_reserve=accumulated_reserve,
        registered_capital=registered_capital,
        ceiling_reached=ceiling_reached,
        ceiling_threshold=ceiling_threshold,
        is_compliant=is_compliant,
        warnings=warnings,
    )


# ═══════════════════════════════════════════════════════════════════════════════
# M6 净利润联动
# ═══════════════════════════════════════════════════════════════════════════════


async def _fetch_m6_net_profit(wp_id: str, db: AsyncSession) -> float | None:
    """从M6未分配利润底稿获取净利润/计提基数。

    尝试从同project下M6底稿的 checklist_responses 读取净利润数据。
    """
    try:
        ctx = await _get_wp_context(wp_id, db)
        project_id = ctx["project_id"]

        result = await db.execute(
            sa.text("""
                SELECT cr.conclusion
                FROM checklist_responses cr
                JOIN working_paper wp ON wp.id = cr.wp_id
                JOIN wp_index wi ON wi.project_id = wp.project_id
                    AND wi.wp_code = 'M6'
                WHERE wp.project_id = :pid
                    AND cr.item_id IN ('M6-net-profit', 'M6-1-net-profit', 'm6:net-profit')
                LIMIT 1
            """),
            {"pid": str(project_id)},
        )
        row = result.fetchone()
        if row and row.conclusion:
            return _parse_num(row.conclusion)
    except Exception as e:  # noqa: BLE001
        logger.warning("M5 计提测试: 从M6获取净利润失败: %s", e)
    return None


# ═══════════════════════════════════════════════════════════════════════════════
# 内部辅助
# ═══════════════════════════════════════════════════════════════════════════════


async def _load_responses(wp_id: str, db: AsyncSession) -> dict[str, dict[str, Any]]:
    """加载底稿的 checklist_responses."""
    result = await db.execute(
        sa.text(
            "SELECT item_id, conclusion, evidence, status "
            "FROM checklist_responses "
            "WHERE wp_id = :wid"
        ),
        {"wid": wp_id},
    )
    responses: dict[str, dict[str, Any]] = {}
    for r in result.fetchall():
        responses[r.item_id] = {
            "conclusion": r.conclusion,
            "evidence": r.evidence,
            "status": r.status,
        }
    return responses

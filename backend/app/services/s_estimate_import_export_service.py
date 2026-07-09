"""S 类计算型底稿 — 导入导出三级业务逻辑服务

提供功能：
- 导出模板 / 导出数据 / 导入数据（xlsx 三级）

适用底稿：
- S21-2 开发支出资本化分析（12月 × 7类目 grid）
- S20-unrelated 与主营业务无关的业务收入明细（动态行）
- S20-noSubstance 不具备商业实质的收入明细（动态行）

多区块分 sheet 导出: S21 按月度网格单 sheet / S20 按两类明细分 sheet
RFC5987 中文文件名: StreamingResponse Content-Disposition header

Requirements: 10.1, 10.2, 10.3, 10.4
"""

from __future__ import annotations

import io
import json
import logging
from typing import Any

from openpyxl import Workbook, load_workbook
from openpyxl.styles import Alignment, Font, PatternFill
from sqlalchemy.ext.asyncio import AsyncSession

logger = logging.getLogger(__name__)

_ROW_LIMIT = 500

# ═══════════════════════════════════════════════════════════════════════════════
# Sheet 配置
# ═══════════════════════════════════════════════════════════════════════════════

# S21-2 成本类目标签
_S21_CATEGORIES = [
    "采购成本",
    "人工成本",
    "脱敏清洗标注整合分析支出",
    "数据权属鉴证费",
    "质量评估费",
    "登记结算费",
    "安全管理费",
]

# S21-2 表头：成本类目 + 1月~12月 + 合计
_S21_2_HEADERS = ["成本类目"] + [f"{m}月" for m in range(1, 13)] + ["合计"]

# S20 明细表头
_S20_DETAIL_HEADERS = ["收入项目名称", "本年度审定数", "上年度审定数", "扣除理由"]
_S20_DETAIL_FIELDS = ["name", "currentYear", "priorYear", "reason"]

_SHEET_CONFIGS: dict[str, dict[str, Any]] = {
    "S21-2": {
        "title": "S21-2 开发支出资本化分析（分月归集）",
        "headers": _S21_2_HEADERS,
        "type": "grid",  # 12月网格特殊处理
    },
    "S20-unrelated": {
        "title": "S20 与主营业务无关的业务收入明细",
        "headers": _S20_DETAIL_HEADERS,
        "fields": _S20_DETAIL_FIELDS,
        "type": "dynamic_rows",
    },
    "S20-noSubstance": {
        "title": "S20 不具备商业实质的收入明细",
        "headers": _S20_DETAIL_HEADERS,
        "fields": _S20_DETAIL_FIELDS,
        "type": "dynamic_rows",
    },
}

_HEADER_FONT = Font(bold=True, size=11, color="FFFFFF")
_HEADER_FILL = PatternFill(start_color="4472C4", end_color="4472C4", fill_type="solid")
_HEADER_ALIGN = Alignment(horizontal="center", vertical="center", wrap_text=True)


# ═══════════════════════════════════════════════════════════════════════════════
# 导出模板
# ═══════════════════════════════════════════════════════════════════════════════


async def export_template(
    wp_id: str,
    db: AsyncSession,
    *,
    sheet: str | None = None,
) -> io.BytesIO:
    """导出空白 xlsx 模板

    Args:
        wp_id: 底稿ID
        db: 数据库会话
        sheet: 可选，指定导出单个 sheet
    """
    wb = Workbook()
    wb.remove(wb.active)  # type: ignore[arg-type]

    sheets_to_export = (
        [sheet] if sheet and sheet in _SHEET_CONFIGS
        else list(_SHEET_CONFIGS.keys())
    )

    for sheet_key in sheets_to_export:
        cfg = _SHEET_CONFIGS[sheet_key]
        ws = wb.create_sheet(title=cfg["title"])

        # 写入表头
        for col_idx, header in enumerate(cfg["headers"], 1):
            cell = ws.cell(row=1, column=col_idx, value=header)
            cell.font = _HEADER_FONT
            cell.fill = _HEADER_FILL
            cell.alignment = _HEADER_ALIGN

        # S21-2 特殊：预填类目名称（行头）
        if cfg["type"] == "grid":
            for row_idx, cat_name in enumerate(_S21_CATEGORIES, 2):
                ws.cell(row=row_idx, column=1, value=cat_name)

        # 设置列宽
        for col_idx in range(1, len(cfg["headers"]) + 1):
            ws.column_dimensions[_col_letter(col_idx)].width = (
                20 if col_idx == 1 else 12
            )

    buffer = io.BytesIO()
    wb.save(buffer)
    buffer.seek(0)
    return buffer


# ═══════════════════════════════════════════════════════════════════════════════
# 导出数据
# ═══════════════════════════════════════════════════════════════════════════════


async def export_data(
    wp_id: str,
    db: AsyncSession,
    *,
    sheet: str | None = None,
) -> io.BytesIO:
    """导出含当前数据的 xlsx

    从 checklist_responses 查询 S 类行数据，写入 xlsx。
    """
    import sqlalchemy as sa
    from app.models.audit_platform_models import ChecklistResponse

    wb = Workbook()
    wb.remove(wb.active)  # type: ignore[arg-type]

    sheets_to_export = (
        [sheet] if sheet and sheet in _SHEET_CONFIGS
        else list(_SHEET_CONFIGS.keys())
    )

    for sheet_key in sheets_to_export:
        cfg = _SHEET_CONFIGS[sheet_key]
        ws = wb.create_sheet(title=cfg["title"])

        # 写入表头
        for col_idx, header in enumerate(cfg["headers"], 1):
            cell = ws.cell(row=1, column=col_idx, value=header)
            cell.font = _HEADER_FONT
            cell.fill = _HEADER_FILL
            cell.alignment = _HEADER_ALIGN

        # 查询 checklist_responses 中的行数据
        item_prefix = f"s-estimate-{sheet_key}-row-"
        stmt = sa.select(ChecklistResponse).where(
            ChecklistResponse.wp_id == wp_id,
            ChecklistResponse.item_id.like(f"{item_prefix}%"),
        ).order_by(ChecklistResponse.item_id)

        result = await db.execute(stmt)
        responses = result.scalars().all()

        if cfg["type"] == "grid":
            # S21-2: 行=类目，列=1-12月+合计
            _export_grid_data(ws, responses)
        else:
            # S20: 动态行明细
            _export_dynamic_rows(ws, responses, cfg["fields"])

        # 设置列宽
        for col_idx in range(1, len(cfg["headers"]) + 1):
            ws.column_dimensions[_col_letter(col_idx)].width = (
                20 if col_idx == 1 else 14
            )

    buffer = io.BytesIO()
    wb.save(buffer)
    buffer.seek(0)
    return buffer


def _export_grid_data(ws: Any, responses: list) -> None:
    """导出 S21-2 网格数据（行=类目，列=1~12月+合计）"""
    # 按类目分组
    category_data: dict[str, list[float]] = {cat: [0.0] * 12 for cat in _S21_CATEGORIES}

    for resp in responses:
        if not resp.remark:
            continue
        try:
            data = json.loads(resp.remark)
        except (json.JSONDecodeError, TypeError):
            continue

        cat_name = data.get("category", "")
        monthly = data.get("monthly", [])
        if cat_name in category_data and isinstance(monthly, list):
            for i, val in enumerate(monthly[:12]):
                try:
                    category_data[cat_name][i] = float(val or 0)
                except (ValueError, TypeError):
                    pass

    # 写入数据行
    for row_idx, cat_name in enumerate(_S21_CATEGORIES, 2):
        ws.cell(row=row_idx, column=1, value=cat_name)
        monthly = category_data[cat_name]
        row_total = 0.0
        for m_idx, val in enumerate(monthly, 2):
            ws.cell(row=row_idx, column=m_idx, value=val)
            row_total += val
        # 合计列
        ws.cell(row=row_idx, column=14, value=row_total)


def _export_dynamic_rows(ws: Any, responses: list, fields: list[str]) -> None:
    """导出 S20 动态行明细"""
    row_num = 2
    for resp in responses:
        if not resp.remark:
            continue
        try:
            data = json.loads(resp.remark)
        except (json.JSONDecodeError, TypeError):
            continue

        for col_idx, field in enumerate(fields, 1):
            value = data.get(field, "")
            ws.cell(row=row_num, column=col_idx, value=value)
        row_num += 1


# ═══════════════════════════════════════════════════════════════════════════════
# 导入数据
# ═══════════════════════════════════════════════════════════════════════════════


async def import_data(
    wp_id: str,
    content: bytes,
    db: AsyncSession,
    *,
    sheet: str | None = None,
) -> dict[str, Any]:
    """导入 xlsx 解析写入 checklist_responses

    解析 xlsx 文件，按 S21-2 / S20 结构写入。

    Returns:
        { imported_count: int, warnings: list[str] }
    """
    import sqlalchemy as sa
    from app.models.audit_platform_models import ChecklistResponse

    target_sheet = sheet or "S21-2"
    if target_sheet not in _SHEET_CONFIGS:
        raise ValueError(f"不支持的 sheet: {target_sheet}")

    cfg = _SHEET_CONFIGS[target_sheet]

    try:
        wb = load_workbook(io.BytesIO(content), read_only=True, data_only=True)
    except Exception as e:
        raise ValueError(f"无法解析 xlsx 文件: {e}")

    # 找到目标工作表
    ws = None
    for ws_name in wb.sheetnames:
        if cfg["title"] in ws_name or target_sheet in ws_name:
            ws = wb[ws_name]
            break
    if ws is None:
        ws = wb.active

    if ws is None:
        raise ValueError("xlsx 文件为空")

    # 验证表头（宽容匹配：只要包含前几个关键列即可）
    header_row = [cell.value for cell in next(ws.iter_rows(min_row=1, max_row=1))]
    expected_first = cfg["headers"][:2]
    missing = [h for h in expected_first if h not in header_row]
    if missing:
        raise ValueError(f"列头不匹配，缺少: {missing}")

    # 分类处理
    if cfg["type"] == "grid":
        result = await _import_grid_data(wp_id, ws, db, target_sheet)
    else:
        result = await _import_dynamic_rows(wp_id, ws, db, target_sheet, cfg["fields"])

    return result


async def _import_grid_data(
    wp_id: str,
    ws: Any,
    db: AsyncSession,
    sheet_key: str,
) -> dict[str, Any]:
    """导入 S21-2 网格数据（按行=类目，列=1~12月）"""
    import sqlalchemy as sa
    from app.models.audit_platform_models import ChecklistResponse

    warnings: list[str] = []
    imported_count = 0

    # 先删除旧数据
    item_prefix = f"s-estimate-{sheet_key}-row-"
    await db.execute(
        sa.delete(ChecklistResponse).where(
            ChecklistResponse.wp_id == wp_id,
            ChecklistResponse.item_id.like(f"{item_prefix}%"),
        )
    )

    for row_idx, row in enumerate(ws.iter_rows(min_row=2, values_only=True), start=1):
        if row_idx > _ROW_LIMIT:
            warnings.append(f"超过最大行限制({_ROW_LIMIT})，后续行已忽略")
            break

        # 跳过全空行
        if all(cell is None or cell == "" for cell in row):
            continue

        cat_name = str(row[0] or "").strip()
        if not cat_name:
            continue

        # 解析 12 月数据（列 2~13）
        monthly: list[float] = []
        for col_idx in range(1, 13):
            val = row[col_idx] if col_idx < len(row) else None
            try:
                monthly.append(float(val or 0))
            except (ValueError, TypeError):
                monthly.append(0.0)

        # 写入 checklist_responses
        item_id = f"{item_prefix}{row_idx:04d}"
        data_json = json.dumps({
            "category": cat_name,
            "monthly": monthly,
        }, ensure_ascii=False)

        db.add(ChecklistResponse(
            wp_id=wp_id,
            item_id=item_id,
            remark=data_json,
        ))
        imported_count += 1

    await db.flush()
    return {"imported_count": imported_count, "warnings": warnings}


async def _import_dynamic_rows(
    wp_id: str,
    ws: Any,
    db: AsyncSession,
    sheet_key: str,
    fields: list[str],
) -> dict[str, Any]:
    """导入 S20 动态行明细"""
    import sqlalchemy as sa
    from app.models.audit_platform_models import ChecklistResponse

    warnings: list[str] = []
    imported_count = 0

    # 先删除旧数据
    item_prefix = f"s-estimate-{sheet_key}-row-"
    await db.execute(
        sa.delete(ChecklistResponse).where(
            ChecklistResponse.wp_id == wp_id,
            ChecklistResponse.item_id.like(f"{item_prefix}%"),
        )
    )

    for row_idx, row in enumerate(ws.iter_rows(min_row=2, values_only=True), start=1):
        if row_idx > _ROW_LIMIT:
            warnings.append(f"超过最大行限制({_ROW_LIMIT})，后续行已忽略")
            break

        # 跳过全空行
        if all(cell is None or cell == "" for cell in row):
            continue

        row_data: dict[str, Any] = {}
        for col_idx, field in enumerate(fields):
            val = row[col_idx] if col_idx < len(row) else None
            if field in ("currentYear", "priorYear"):
                try:
                    row_data[field] = float(val or 0)
                except (ValueError, TypeError):
                    row_data[field] = 0.0
            else:
                row_data[field] = str(val) if val is not None else ""

        # 写入 checklist_responses
        item_id = f"{item_prefix}{row_idx:04d}"
        data_json = json.dumps(row_data, ensure_ascii=False)

        db.add(ChecklistResponse(
            wp_id=wp_id,
            item_id=item_id,
            remark=data_json,
        ))
        imported_count += 1

    await db.flush()
    return {"imported_count": imported_count, "warnings": warnings}


# ═══════════════════════════════════════════════════════════════════════════════
# 辅助工具
# ═══════════════════════════════════════════════════════════════════════════════


def get_sheet_filename(sheet_key: str, suffix: str) -> str:
    """获取导出文件名（中文，用于 RFC5987 编码）

    Args:
        sheet_key: sheet 标识（S21-2 / S20-unrelated / S20-noSubstance）
        suffix: 后缀（模板/数据）
    """
    name_map = {
        "S21-2": "S21-2 开发支出资本化分析",
        "S20-unrelated": "S20 与主营无关收入明细",
        "S20-noSubstance": "S20 不具备商业实质收入明细",
    }
    base = name_map.get(sheet_key, f"S类底稿-{sheet_key}")
    return f"{base}_{suffix}.xlsx"


def _col_letter(col_idx: int) -> str:
    """列号转列字母（1→A, 2→B, ...26→Z, 27→AA）"""
    result = ""
    while col_idx > 0:
        col_idx, remainder = divmod(col_idx - 1, 26)
        result = chr(65 + remainder) + result
    return result

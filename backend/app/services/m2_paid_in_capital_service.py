"""M2 实收资本（股本）— 服务层

外币折算 + 验资核对 + 按出资人汇总 + 导入导出

科目编码: 4001实收资本/股本（贷方/权益类！）
公式方向: 期末=期初+贷方-借方（增资在贷方，减资在借方）

Requirements: 4.2, 5.2, 7.7
"""

from __future__ import annotations

import io
import logging
from typing import Any

from sqlalchemy.ext.asyncio import AsyncSession
import sqlalchemy as sa

logger = logging.getLogger(__name__)

# 导入导出涉及的3个sheet（动态行表格）
EXPORT_SHEETS = ["M2-2-listed", "M2-2-unlisted", "M2-4"]
SHEET_NAMES = {
    "M2-2-listed": "M2-2 明细表(上市)",
    "M2-2-unlisted": "M2-2 明细表(非上市)",
    "M2-4": "M2-4 外币投资汇率",
}


# ═══════════════════════════════════════════════════════════════════════════════
# 纯函数：外币折算（Requirements: 4.2）
# ═══════════════════════════════════════════════════════════════════════════════


def calculate_fx_conversion(amount: float, rate: float) -> float:
    """折算本位币 = 原币出资 × 出资日汇率"""
    return round(amount * rate, 2)


def calculate_fx_diff(converted: float, booked: float) -> float:
    """折算差异 = 折算本位币 - 账面本位币"""
    return round(converted - booked, 2)


# ═══════════════════════════════════════════════════════════════════════════════
# 纯函数：验资核对（Requirements: 5.2）
# ═══════════════════════════════════════════════════════════════════════════════


def calculate_verify_diff(paid: float, verified: float) -> float:
    """验资差异 = 实缴出资 - 验资金额"""
    return round(paid - verified, 2)


def calculate_paid_in_rate(paid: float, subscribed: float) -> float:
    """出资到位率 = 实缴出资 / 认缴出资（guard div0）"""
    if subscribed == 0:
        return 0.0
    return round(paid / subscribed, 6)


# ═══════════════════════════════════════════════════════════════════════════════
# 权益类公式验证
# ═══════════════════════════════════════════════════════════════════════════════


def validate_equity_formulas(
    begin: float,
    credit: float,
    debit: float,
    reported_end: float | None = None,
) -> dict[str, Any]:
    """校验权益类公式方向：期末=期初+贷方-借方

    实收资本为贷方/权益类科目：
    - 增资（含验资确认）在贷方增加
    - 减资（减少注册资本）在借方减少
    """
    calculated_end = begin + credit - debit

    result: dict[str, Any] = {
        "begin": begin,
        "credit": credit,
        "debit": debit,
        "calculatedEnd": round(calculated_end, 2),
        "formula": "期末 = 期初 + 贷方(增资) - 借方(减资)",
        "direction": "credit",  # 贷方/权益类
        "isValid": True,
    }

    if reported_end is not None:
        diff = round(calculated_end - reported_end, 2)
        result["reportedEnd"] = reported_end
        result["diff"] = diff
        result["isValid"] = abs(diff) < 0.01

    return result


# ═══════════════════════════════════════════════════════════════════════════════
# 按出资人汇总（Requirements: 5.2）
# ═══════════════════════════════════════════════════════════════════════════════


async def summarize_by_investor(
    wp_id: str,
    db: AsyncSession,
) -> dict[str, Any]:
    """从 checklist_responses 按出资人汇总实收资本数据

    Returns:
        {
            investors: list[{name, subscribed, paid, verified, paid_in_rate, verify_diff}],
            total_subscribed: float,
            total_paid: float,
            total_verified: float,
            overall_paid_in_rate: float,
        }
    """
    result = await db.execute(
        sa.text("""
            SELECT item_id, value
            FROM checklist_responses
            WHERE wp_id = :wp_id AND item_id LIKE 'M2-5-%'
            ORDER BY item_id
        """),
        {"wp_id": wp_id},
    )
    rows = result.fetchall()

    # 按出资人行分组解析
    row_data: dict[str, dict[str, Any]] = {}
    for row in rows:
        item_id = row.item_id
        value = row.value
        # item_id format: M2-5-{row_idx}-{field}
        parts = item_id.split("-")
        if len(parts) >= 4:
            row_key = parts[2]
            field = "-".join(parts[3:])
            if row_key not in row_data:
                row_data[row_key] = {}
            row_data[row_key][field] = value

    investors: list[dict[str, Any]] = []
    total_subscribed = 0.0
    total_paid = 0.0
    total_verified = 0.0

    for row_key in sorted(row_data.keys()):
        data = row_data[row_key]
        name = data.get("investor", "")
        subscribed = _safe_float(data.get("subscribed", 0))
        paid = _safe_float(data.get("paid", 0))
        verified = _safe_float(data.get("verified", 0))

        verify_diff = calculate_verify_diff(paid, verified)
        paid_in_rate = calculate_paid_in_rate(paid, subscribed)

        investors.append({
            "name": name,
            "subscribed": subscribed,
            "paid": paid,
            "verified": verified,
            "paid_in_rate": paid_in_rate,
            "verify_diff": verify_diff,
        })

        total_subscribed += subscribed
        total_paid += paid
        total_verified += verified

    overall_rate = calculate_paid_in_rate(total_paid, total_subscribed)

    return {
        "investors": investors,
        "total_subscribed": round(total_subscribed, 2),
        "total_paid": round(total_paid, 2),
        "total_verified": round(total_verified, 2),
        "overall_paid_in_rate": overall_rate,
    }


# ═══════════════════════════════════════════════════════════════════════════════
# 外币折算批量计算
# ═══════════════════════════════════════════════════════════════════════════════


async def calc_fx_batch(
    db: AsyncSession,
    wp_id: str,
) -> dict[str, Any]:
    """从 checklist_responses 取 M2-4 数据并批量计算外币折算

    折算本位币 = 原币出资 × 出资日汇率
    折算差异 = 折算本位币 - 账面本位币
    """
    result = await db.execute(
        sa.text("""
            SELECT item_id, value
            FROM checklist_responses
            WHERE wp_id = :wp_id AND item_id LIKE 'M2-4-%'
            ORDER BY item_id
        """),
        {"wp_id": wp_id},
    )
    rows = result.fetchall()

    row_data: dict[str, dict[str, Any]] = {}
    for row in rows:
        item_id = row.item_id
        value = row.value
        parts = item_id.split("-")
        if len(parts) >= 4:
            row_key = parts[2]
            field = "-".join(parts[3:])
            if row_key not in row_data:
                row_data[row_key] = {}
            row_data[row_key][field] = value

    items: list[dict[str, Any]] = []
    total_fx_diff = 0.0
    material_threshold = 1000.0

    for row_key in sorted(row_data.keys()):
        data = row_data[row_key]
        original_amount = _safe_float(data.get("original_amount", 0))
        rate = _safe_float(data.get("rate", 0))
        booked = _safe_float(data.get("booked", 0))

        converted = calculate_fx_conversion(original_amount, rate)
        fx_diff = calculate_fx_diff(converted, booked)
        total_fx_diff += fx_diff

        items.append({
            "investor": data.get("investor", ""),
            "currency": data.get("currency", ""),
            "original_amount": original_amount,
            "rate": rate,
            "converted": converted,
            "booked": booked,
            "fx_diff": fx_diff,
        })

    return {
        "items": items,
        "total_fx_diff": round(total_fx_diff, 2),
        "has_material_diff": abs(total_fx_diff) > material_threshold,
    }


# ═══════════════════════════════════════════════════════════════════════════════
# 验资核对批量计算
# ═══════════════════════════════════════════════════════════════════════════════


async def calc_verify_batch(
    db: AsyncSession,
    wp_id: str,
) -> dict[str, Any]:
    """从 checklist_responses 取 M2-5 数据并批量计算验资差异

    验资差异 = 实缴出资 - 验资金额
    出资到位率 = 实缴出资 / 认缴出资
    """
    result = await db.execute(
        sa.text("""
            SELECT item_id, value
            FROM checklist_responses
            WHERE wp_id = :wp_id AND item_id LIKE 'M2-5-%'
            ORDER BY item_id
        """),
        {"wp_id": wp_id},
    )
    rows = result.fetchall()

    row_data: dict[str, dict[str, Any]] = {}
    for row in rows:
        item_id = row.item_id
        value = row.value
        parts = item_id.split("-")
        if len(parts) >= 4:
            row_key = parts[2]
            field = "-".join(parts[3:])
            if row_key not in row_data:
                row_data[row_key] = {}
            row_data[row_key][field] = value

    items: list[dict[str, Any]] = []
    total_paid = 0.0
    total_verified = 0.0
    material_threshold = 1000.0

    for row_key in sorted(row_data.keys()):
        data = row_data[row_key]
        investor = data.get("investor", "")
        subscribed = _safe_float(data.get("subscribed", 0))
        paid = _safe_float(data.get("paid", 0))
        verified = _safe_float(data.get("verified", 0))

        verify_diff = calculate_verify_diff(paid, verified)
        paid_in_rate = calculate_paid_in_rate(paid, subscribed)
        total_paid += paid
        total_verified += verified

        items.append({
            "investor": investor,
            "subscribed": subscribed,
            "paid": paid,
            "verified": verified,
            "verify_diff": verify_diff,
            "paid_in_rate": paid_in_rate,
        })

    total_diff = round(total_paid - total_verified, 2)

    return {
        "items": items,
        "total_paid": round(total_paid, 2),
        "total_verified": round(total_verified, 2),
        "total_diff": total_diff,
        "has_material_diff": abs(total_diff) > material_threshold,
    }


# ═══════════════════════════════════════════════════════════════════════════════
# 导出模板
# ═══════════════════════════════════════════════════════════════════════════════


async def export_template(wp_id: str, db: AsyncSession, sheet: str | None = None) -> io.BytesIO:
    """导出空白xlsx模板（3个sheet：M2-2上市明细/M2-2非上市明细/M2-4外币投资）"""
    from openpyxl import Workbook

    wb = Workbook()
    wb.remove(wb.active)

    sheets_to_export = [sheet] if sheet and sheet in EXPORT_SHEETS else EXPORT_SHEETS

    if "M2-2-listed" in sheets_to_export:
        ws = wb.create_sheet("M2-2 明细表(上市)")
        ws.append([
            "股东名称", "股份性质", "期初股数", "期初金额",
            "本期增加股数", "本期增加金额", "本期减少股数", "本期减少金额",
            "期末股数", "期末金额", "持股比例(%)", "备注",
        ])

    if "M2-2-unlisted" in sheets_to_export:
        ws = wb.create_sheet("M2-2 明细表(非上市)")
        ws.append([
            "出资人", "出资方式", "认缴出资", "实缴出资",
            "期初金额", "本期增资", "本期减资", "期末金额",
            "出资比例(%)", "出资日期", "备注",
        ])

    if "M2-4" in sheets_to_export:
        ws = wb.create_sheet("M2-4 外币投资汇率")
        ws.append([
            "出资人", "币种", "原币出资", "出资日汇率",
            "折算本位币", "账面本位币", "折算差异",
        ])

    buffer = io.BytesIO()
    wb.save(buffer)
    buffer.seek(0)
    return buffer


# ═══════════════════════════════════════════════════════════════════════════════
# 导出数据
# ═══════════════════════════════════════════════════════════════════════════════


async def export_data(wp_id: str, db: AsyncSession, sheet: str | None = None) -> io.BytesIO:
    """导出含当前数据的xlsx（3个sheet）"""
    from openpyxl import Workbook

    wb = Workbook()
    wb.remove(wb.active)

    sheets_to_export = [sheet] if sheet and sheet in EXPORT_SHEETS else EXPORT_SHEETS

    # 读取 checklist_responses 中的 M2 数据
    like_patterns = []
    if "M2-2-listed" in sheets_to_export:
        like_patterns.append("M2-2-listed-%")
    if "M2-2-unlisted" in sheets_to_export:
        like_patterns.append("M2-2-unlisted-%")
    if "M2-4" in sheets_to_export:
        like_patterns.append("M2-4-%")

    all_rows: list[Any] = []
    for pattern in like_patterns:
        result = await db.execute(
            sa.text("""
                SELECT item_id, value
                FROM checklist_responses
                WHERE wp_id = :wp_id AND item_id LIKE :pattern
                ORDER BY item_id
            """),
            {"wp_id": wp_id, "pattern": pattern},
        )
        all_rows.extend(result.fetchall())

    # 分组
    sheet_data: dict[str, dict[str, dict[str, Any]]] = {
        "M2-2-listed": {},
        "M2-2-unlisted": {},
        "M2-4": {},
    }

    for row in all_rows:
        item_id = row.item_id
        value = row.value
        # M2-2-listed-001-field or M2-2-unlisted-001-field or M2-4-001-field
        if item_id.startswith("M2-2-listed-"):
            rest = item_id[len("M2-2-listed-"):]
            parts = rest.split("-", 1)
            sheet_key = "M2-2-listed"
        elif item_id.startswith("M2-2-unlisted-"):
            rest = item_id[len("M2-2-unlisted-"):]
            parts = rest.split("-", 1)
            sheet_key = "M2-2-unlisted"
        elif item_id.startswith("M2-4-"):
            rest = item_id[len("M2-4-"):]
            parts = rest.split("-", 1)
            sheet_key = "M2-4"
        else:
            continue

        if len(parts) >= 2:
            row_key = parts[0]
            field = parts[1]
            if row_key not in sheet_data[sheet_key]:
                sheet_data[sheet_key][row_key] = {}
            sheet_data[sheet_key][row_key][field] = value

    # M2-2 上市明细
    if "M2-2-listed" in sheets_to_export:
        ws = wb.create_sheet("M2-2 明细表(上市)")
        ws.append(["股东名称", "股份性质", "期初股数", "期初金额",
                   "本期增加股数", "本期增加金额", "本期减少股数", "本期减少金额",
                   "期末股数", "期末金额", "持股比例(%)", "备注"])
        for rk in sorted(sheet_data["M2-2-listed"].keys()):
            d = sheet_data["M2-2-listed"][rk]
            ws.append([
                d.get("shareholder", ""), d.get("share_type", ""),
                d.get("begin_shares", ""), d.get("begin_amount", ""),
                d.get("increase_shares", ""), d.get("increase_amount", ""),
                d.get("decrease_shares", ""), d.get("decrease_amount", ""),
                d.get("end_shares", ""), d.get("end_amount", ""),
                d.get("ratio", ""), d.get("note", ""),
            ])

    # M2-2 非上市明细
    if "M2-2-unlisted" in sheets_to_export:
        ws = wb.create_sheet("M2-2 明细表(非上市)")
        ws.append(["出资人", "出资方式", "认缴出资", "实缴出资",
                   "期初金额", "本期增资", "本期减资", "期末金额",
                   "出资比例(%)", "出资日期", "备注"])
        for rk in sorted(sheet_data["M2-2-unlisted"].keys()):
            d = sheet_data["M2-2-unlisted"][rk]
            ws.append([
                d.get("investor", ""), d.get("invest_method", ""),
                d.get("subscribed", ""), d.get("paid", ""),
                d.get("begin_amount", ""), d.get("increase", ""),
                d.get("decrease", ""), d.get("end_amount", ""),
                d.get("ratio", ""), d.get("invest_date", ""),
                d.get("note", ""),
            ])

    # M2-4 外币投资
    if "M2-4" in sheets_to_export:
        ws = wb.create_sheet("M2-4 外币投资汇率")
        ws.append(["出资人", "币种", "原币出资", "出资日汇率",
                   "折算本位币", "账面本位币", "折算差异"])
        for rk in sorted(sheet_data["M2-4"].keys()):
            d = sheet_data["M2-4"][rk]
            ws.append([
                d.get("investor", ""), d.get("currency", ""),
                d.get("original_amount", ""), d.get("rate", ""),
                d.get("converted", ""), d.get("booked", ""),
                d.get("fx_diff", ""),
            ])

    buffer = io.BytesIO()
    wb.save(buffer)
    buffer.seek(0)
    return buffer


# ═══════════════════════════════════════════════════════════════════════════════
# 导入数据
# ═══════════════════════════════════════════════════════════════════════════════


async def import_data(
    wp_id: str,
    content: bytes,
    db: AsyncSession,
    sheet: str | None = None,
) -> dict[str, Any]:
    """导入xlsx解析写入checklist_responses（支持3个sheet）"""
    from openpyxl import load_workbook

    buffer = io.BytesIO(content)
    try:
        wb = load_workbook(buffer, read_only=True, data_only=True)
    except Exception as e:
        raise ValueError(f"无法解析xlsx文件: {e}")

    imported_count = 0
    warnings: list[str] = []

    # sheet名称匹配（容错）
    sheet_mapping: dict[str, list[str]] = {
        "M2-2-listed": ["M2-2 明细表(上市)", "M2-2 明细表（上市）", "M2-2-listed", "明细表(上市)"],
        "M2-2-unlisted": ["M2-2 明细表(非上市)", "M2-2 明细表（非上市）", "M2-2-unlisted", "明细表(非上市)"],
        "M2-4": ["M2-4 外币投资汇率", "M2-4", "外币投资汇率"],
    }

    sheets_to_import = {sheet: sheet_mapping[sheet]} if sheet and sheet in sheet_mapping else sheet_mapping

    for sheet_key, possible_names in sheets_to_import.items():
        ws = None
        for name in possible_names:
            if name in wb.sheetnames:
                ws = wb[name]
                break

        if ws is None:
            continue

        row_idx = 0
        for row in ws.iter_rows(min_row=2, values_only=True):
            if not row or all(cell is None or cell == "" for cell in row):
                continue

            row_idx += 1
            fields = _get_field_mapping(sheet_key)

            for col_idx, field in enumerate(fields):
                if col_idx < len(row) and row[col_idx] is not None:
                    item_id = f"{sheet_key}-{row_idx:03d}-{field}"
                    value = str(row[col_idx])

                    await db.execute(
                        sa.text("""
                            INSERT INTO checklist_responses (wp_id, item_id, value)
                            VALUES (:wp_id, :item_id, :value)
                            ON CONFLICT (wp_id, item_id) DO UPDATE SET value = :value
                        """),
                        {"wp_id": wp_id, "item_id": item_id, "value": value},
                    )
                    imported_count += 1

    wb.close()
    return {"imported_count": imported_count, "warnings": warnings}


# ═══════════════════════════════════════════════════════════════════════════════
# 辅助函数
# ═══════════════════════════════════════════════════════════════════════════════


def _safe_float(val: Any) -> float:
    """安全浮点转换"""
    try:
        return float(val)
    except (TypeError, ValueError):
        return 0.0


def _get_field_mapping(sheet_key: str) -> list[str]:
    """获取sheet列→field名映射"""
    if sheet_key == "M2-2-listed":
        return [
            "shareholder", "share_type", "begin_shares", "begin_amount",
            "increase_shares", "increase_amount", "decrease_shares", "decrease_amount",
            "end_shares", "end_amount", "ratio", "note",
        ]
    elif sheet_key == "M2-2-unlisted":
        return [
            "investor", "invest_method", "subscribed", "paid",
            "begin_amount", "increase", "decrease", "end_amount",
            "ratio", "invest_date", "note",
        ]
    elif sheet_key == "M2-4":
        return [
            "investor", "currency", "original_amount", "rate",
            "converted", "booked", "fx_diff",
        ]
    return []

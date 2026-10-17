"""财务报表模板行绑定：按「模板行名」把库内报表行绑定到模板占位。

为什么需要（2026-10-16 真实数据实测，spec ``deliverable-trio-generation-accuracy``）
--------------------------------------------------------------------------------
- 报表模板 xlsx 的 ``{{row:CODE}}`` / ``{{eq:CODE}}`` / ``{{imp:CODE}}`` 编码出自
  ``report_config_seed.json``，与模板逐格一致；
- 库内 ``report_config`` 却是两套编号的混合体（soe_standalone 226 个占位里 113 个
  同码不同名，listed 64/146），旧导出按编码直取 ⇒ 重药安徽资产负债表「衍生金融资产」
  拿到应收账款 5.96 亿、资产总计 31.8 亿（报表模块为 11.5 亿）；
- 平台约定报表编号以库为准（``scripts/fix/remap_note_report_row_codes.py``），本模块
  **不改库编号**，只在导出 / 差异检测读取时按行名绑定。

绑定规则（每个模板「行」一个出现位，键 = sheet!行号!编码）
-------------------------------------------------------
1. 候选 = 同编码族（BS/IS/CFS/CFSS/EQ/IMP）里行名规范化 key 与标记（△▲#*）都相同的库内行；
2. 模板里该 key 只出现一次：按 (非空壳行, 被报表合计公式引用, 同编码, 有公式, 有金额, 编码序) 取优 ——
   库内同名行常有通用编号 / 模板编号两份且公式不同（安徽「其他应付款」一份含应付利息），
   报表模块的合计只引用其中一份，取这份导出合计才与报表模块一致；
3. 模板里该 key 出现多次：先按同编码配对（同一编码在模板多处出现时，前后编码连号的那处优先），
   余下按出现顺序 ↔ 编码顺序配对；
4. 无候选 → 不取数（同编码但行名不同记为 ``conflicts``，绝不按编码硬取）；
5. 模板行本身没有行名（合成模板等）→ 无从核对行名，按编码取（与旧口径一致）。
"""

from __future__ import annotations

import logging
import re
import unicodedata
from collections import Counter, defaultdict
from dataclasses import dataclass, field
from decimal import Decimal, InvalidOperation
from functools import lru_cache
from typing import Any, Iterable

logger = logging.getLogger(__name__)

__all__ = [
    "SLOT_KEYED_MARKER",
    "BindingResult",
    "TemplateRowSlot",
    "TemplateSlots",
    "bind_rows_to_template",
    "code_family",
    "expected_for_cell",
    "load_template_slots",
    "normalize_row_label",
    "slot_key",
]

# 整格占位：{{row:BS-002:current}} / {{row:BS-002:prior:parent}} /
# {{eq:EQ-001:current_year:share_capital}} / {{imp:IMP-001:opening_balance}}
_PLACEHOLDER_RE = re.compile(r"^\{\{(row|eq|imp):([A-Za-z]+-[^:}]+):[^}]*\}\}$")
_MARKER_RE = re.compile(r"^[△▲#*]+")
_ENUM_RE = re.compile(
    r"^(?:[一二三四五六七八九十]+[、.]"  # 一、 / 一.
    r"|\([一二三四五六七八九十0-9]+\)"  # (一) / (1)（NFKC 后全角括号已转半角）
    r"|[0-9]+[、.](?!\d)"  # 1、 / 1.（排除 1.5 这类小数）
    r"|[①-⑳])"
)
_LEAD_WORD_RE = re.compile(r"^(?:其中|加|减)[:]")
_QUOTE_RE = re.compile(r"[“”\"'‘’]")
_DASH_RE = re.compile(r"[—–−]")
_TAIL_HINT_RE = re.compile(r"\((?:[^()]*填列[^()]*|或股本|或所有者权益|或股东权益)\)$")
# 同义行名（国企模板「年」口径 vs 库内通用「期」口径等），规范化后整键替换
_KEY_SYNONYMS = {
    "上期期末余额": "上年年末余额",
    "本期期初余额": "本年年初余额",
    "本期期末余额": "本年年末余额",
    "本期增减变动金额": "本年增减变动金额",
}


def code_family(code: str | None) -> str:
    """编码族：``BS-002`` → ``BS``；``CFSS-001`` → ``CFSS``；无前缀返回空串。"""
    m = re.match(r"^([A-Za-z]+)-", str(code or ""))
    return m.group(1).upper() if m else ""


def normalize_row_label(text: Any) -> tuple[str, str]:
    """报表行名规范化 → ``(key, markers)``。

    - NFKC 统一全半角（括号 / 冒号 / 数字 / 破折号），去全部空白（含全角空格）；
    - 反复剥离开头的 △▲#* 标记（单独返回）、序号（一、/（一）/1./①）、「其中：/加：/减：」；
    - 去引号，破折号归一，去结尾「（…填列）」「（或股本）」类括注。
    """
    s = unicodedata.normalize("NFKC", str(text or ""))
    s = re.sub(r"\s+", "", s)
    markers: set[str] = set()
    changed = True
    while changed and s:
        changed = False
        m = _MARKER_RE.match(s)
        if m:
            markers.update(m.group(0))
            s = s[m.end():]
            changed = True
            continue
        m = _ENUM_RE.match(s)
        if m:
            s = s[m.end():]
            changed = True
            continue
        m = _LEAD_WORD_RE.match(s)
        if m:
            s = s[m.end():]
            changed = True
    s = _QUOTE_RE.sub("", s)
    s = _DASH_RE.sub("-", s)
    s = _TAIL_HINT_RE.sub("", s)
    s = s.rstrip(":")
    s = s.replace("以及", "及")
    s = _KEY_SYNONYMS.get(s, s)
    return s, "".join(sorted(markers))


def slot_key(sheet: str, row: int, code: str) -> str:
    """模板出现位键：``sheet!行号!编码``（同一编码可在模板多处出现且行名不同）。"""
    return f"{sheet}!{row}!{code}"


@dataclass(frozen=True)
class TemplateRowSlot:
    """模板中一个「行 × 编码」出现位。"""

    sheet: str
    row: int
    code: str
    family: str
    kind: str  # row / eq / imp
    label: str
    key: str
    markers: str
    order: int  # 全模板出现序（sheet 序 × 行序）

    @property
    def slot_key(self) -> str:
        return slot_key(self.sheet, self.row, self.code)


@dataclass
class TemplateSlots:
    """一份模板的全部出现位 + 「模板自算行」（金额格为 Excel 公式的行）行名键。"""

    slots: list[TemplateRowSlot] = field(default_factory=list)
    #: (family, key, markers) —— 模板用公式自算的行（合计 / 小计），库内同名行不算「遗漏」
    formula_keys: set[tuple[str, str, str]] = field(default_factory=set)
    #: sheet → 该 sheet 的主编码族（按占位多数）
    sheet_family: dict[str, str] = field(default_factory=dict)
    #: (sheet, 行号) → 模板行名原文（含公式行），供导出后校验按行名找格
    row_labels: dict[tuple[str, int], str] = field(default_factory=dict)


def _is_placeholder(value: Any) -> bool:
    return isinstance(value, str) and "{{" in value


def _is_formula(value: Any) -> bool:
    return isinstance(value, str) and value.startswith("=")


def _row_label(cells: list[Any]) -> str | None:
    """同行最左侧的文字格（排除占位 / 公式 / 纯数字 / 「——」）。"""
    for c in cells:
        v = getattr(c, "value", None)
        if not isinstance(v, str):
            continue
        t = v.strip()
        if not t or _is_placeholder(t) or _is_formula(t):
            continue
        if re.fullmatch(r"[\d.\-—]+", t):
            continue
        return t
    return None


def scan_workbook_slots(wb) -> TemplateSlots:
    """扫描已加载的模板工作簿（未填充）得到出现位。"""
    out = TemplateSlots()
    order = 0
    for ws in wb.worksheets:
        fam_counter: Counter[str] = Counter()
        formula_rows: list[tuple[int, str]] = []
        for row in ws.iter_rows():
            cells = list(row)
            if not cells:
                continue
            label = _row_label(cells)
            row_idx = getattr(cells[0], "row", None)
            if row_idx is None:
                continue
            if label:
                out.row_labels[(ws.title, row_idx)] = label
            seen_codes: set[str] = set()
            has_formula_amount = False
            for c in cells:
                v = getattr(c, "value", None)
                if _is_formula(v):
                    has_formula_amount = True
                    continue
                if not isinstance(v, str):
                    continue
                m = _PLACEHOLDER_RE.match(v.strip())
                if not m:
                    continue
                kind, code = m.group(1), m.group(2)
                if code in seen_codes:
                    continue
                seen_codes.add(code)
                fam = code_family(code)
                fam_counter[fam] += 1
                key, markers = normalize_row_label(label or "")
                out.slots.append(
                    TemplateRowSlot(
                        sheet=ws.title,
                        row=row_idx,
                        code=code,
                        family=fam,
                        kind=kind,
                        label=label or "",
                        key=key,
                        markers=markers,
                        order=order,
                    )
                )
                order += 1
            if has_formula_amount and label and not seen_codes:
                formula_rows.append((row_idx, label))
        sheet_fam = fam_counter.most_common(1)[0][0] if fam_counter else ""
        out.sheet_family[ws.title] = sheet_fam
        for _r, label in formula_rows:
            key, markers = normalize_row_label(label)
            out.formula_keys.add((sheet_fam, key, markers))
    return out


@lru_cache(maxsize=16)
def _load_slots_cached(path: str, mtime: float) -> TemplateSlots:  # noqa: ARG001 — mtime 作缓存键
    from openpyxl import load_workbook

    wb = load_workbook(path, read_only=True, data_only=False)
    try:
        return scan_workbook_slots(wb)
    finally:
        wb.close()


def load_template_slots(template_key: str) -> TemplateSlots | None:
    """按变体键读模板出现位（文件改动即失效）；模板缺失返回 ``None``。

    模板路径经 ``report_excel_exporter.resolve_fs_template_path``（报表域里唯一经 manifest 定位模板的地方）。
    """
    from app.services.report_excel_exporter import resolve_fs_template_path

    p = resolve_fs_template_path(template_key)
    if p is None:
        return None
    return _load_slots_cached(str(p), p.stat().st_mtime)


# ---------------------------------------------------------------------------
# 绑定
# ---------------------------------------------------------------------------


def _amount(v: Any) -> Decimal | None:
    if v is None or v == "":
        return None
    try:
        return Decimal(str(v))
    except (InvalidOperation, ValueError):
        return None


def _nonzero(row: dict) -> bool:
    for k in ("current_period_amount", "prior_period_amount"):
        a = _amount(row.get(k))
        if a is not None and a != 0:
            return True
    return False


def _has_formula(row: dict) -> bool:
    return bool(str(row.get("formula_used") or "").strip())


def _is_weak(row: dict) -> bool:
    """无公式且无金额的行（多为 seed 同步时只建了结构的占位行）。"""
    return not _has_formula(row) and not _nonzero(row)


@dataclass
class BindingResult:
    """绑定结果：按出现位 / 按模板编码两种索引 + 诊断信息。"""

    by_slot: dict[str, dict] = field(default_factory=dict)
    by_code: dict[str, dict] = field(default_factory=dict)
    unresolved: list[TemplateRowSlot] = field(default_factory=list)
    #: (出现位, 库内同编码行的行名) —— 旧导出会把这些行的金额写错格
    conflicts: list[tuple[TemplateRowSlot, str]] = field(default_factory=list)
    #: 库内有金额、但模板既无对应占位也无同名自算行的报表行
    unbound_nonzero: list[dict] = field(default_factory=list)
    methods: Counter = field(default_factory=Counter)

    def lookup(self, sheet: str, row: int, code: str) -> dict | None:
        hit = self.by_slot.get(slot_key(sheet, row, code))
        if hit is not None:
            return hit
        return None

    def amounts_lookup(self) -> dict[str, Any]:
        """差异检测口径：出现位键 + 模板编码键 → 金额（出现位键优先），带 ``SLOT_KEYED_MARKER``（见 ``expected_for_cell``）。"""
        out: dict[str, Any] = {}
        for code, r in self.by_code.items():
            out[code] = {
                "current_period_amount": r.get("current_period_amount"),
                "prior_period_amount": r.get("prior_period_amount"),
            }
        for k, r in self.by_slot.items():
            out[k] = {
                "current_period_amount": r.get("current_period_amount"),
                "prior_period_amount": r.get("prior_period_amount"),
            }
        out[SLOT_KEYED_MARKER] = True
        return out


#: 差异检测权威值按模板出现位（sheet!行!编码）键入的标记键
SLOT_KEYED_MARKER = "__slot_keyed__"


def expected_for_cell(expected: dict[str, Any], sheet: str, coord: str | None, row_code: str) -> dict[str, Any] | None:
    """取某格的重算权威值：出现位键入时只认出现位（未绑定的格不回退编码，免误报），否则按编码。

    🔴 库内 report_config 与模板占位是两套编号，导出按模板行名绑定（``bind_rows_to_template``）；差异检测
    必须用同一绑定，否则会把「导出正确、编码不同」的格全报成差异 → confirm 闸门误拦（2026-10-16）。
    """
    if expected.get(SLOT_KEYED_MARKER):
        m = re.match(r"^[A-Z]+(\d+)$", coord or "")
        return expected.get(slot_key(sheet, int(m.group(1)), row_code)) if m else None
    return expected.get(row_code)


_ROW_REF_RE = re.compile(r"\bROW\(\s*'([A-Za-z]+-[^']+)'\s*\)")
_SUM_ROW_RE = re.compile(r"\bSUM_ROW\(\s*'([A-Za-z]+-[^']+)'\s*,\s*'([A-Za-z]+-[^']+)'\s*\)")


def referenced_row_codes(rows: Iterable[dict]) -> set[str]:
    """被其它行合计公式（``ROW('X')`` / ``SUM_ROW('A','B')``）引用的行编码。

    库内同名行常有两份（通用编号 + 模板编号，公式可能不同，如「其他应付款」一份含应付利息
    一份不含）；报表模块自己的合计只引用其中一份 —— 绑定取这份，导出合计才与报表模块一致。
    """
    rows = list(rows)
    all_codes = sorted({str(r.get("row_code")) for r in rows if r.get("row_code")})
    refs: set[str] = set()
    for r in rows:
        f = str(r.get("formula_used") or "")
        if not f:
            continue
        refs.update(_ROW_REF_RE.findall(f))
        for a, b in _SUM_ROW_RE.findall(f):
            fam = code_family(a)
            refs.update(c for c in all_codes if code_family(c) == fam and a <= c <= b)
    return refs


def _rank(candidates: list[dict], code: str, referenced: set[str] | frozenset[str] = frozenset()) -> dict:
    return sorted(
        candidates,
        key=lambda r: (
            _is_weak(r),
            str(r.get("row_code") or "") not in referenced,
            r.get("row_code") != code,
            not _has_formula(r),
            not _nonzero(r),
            str(r.get("row_code") or ""),
        ),
    )[0]


def _code_neighbor_score(slot: TemplateRowSlot, ordered: list[TemplateRowSlot]) -> int:
    """同一编码在模板多处出现时，前后出现位编码连号（code-1 / code+1）的那处是本编码的「本位」。"""
    m = re.match(r"^([A-Za-z]+-)(\d+)$", slot.code)
    if not m:
        return 0
    prefix, num = m.group(1), int(m.group(2))
    width = len(m.group(2))
    same_sheet = [s for s in ordered if s.sheet == slot.sheet]
    try:
        i = next(k for k, s in enumerate(same_sheet) if s.slot_key == slot.slot_key)
    except StopIteration:
        return 0
    score = 0
    if i > 0 and same_sheet[i - 1].code == f"{prefix}{num - 1:0{width}d}":
        score += 1
    if i + 1 < len(same_sheet) and same_sheet[i + 1].code == f"{prefix}{num + 1:0{width}d}":
        score += 1
    return score


def bind_rows_to_template(rows: Iterable[dict], slots: TemplateSlots) -> BindingResult:
    """把库内报表行（``_load_report_data`` 的行 dict）绑定到模板出现位。"""
    result = BindingResult()
    db_rows = [r for r in rows if isinstance(r, dict) and r.get("row_code")]
    by_code: dict[str, dict] = {}
    groups: dict[tuple[str, str, str], list[dict]] = defaultdict(list)
    for r in sorted(db_rows, key=lambda x: str(x.get("row_code"))):
        code = str(r["row_code"])
        by_code.setdefault(code, r)
        key, markers = normalize_row_label(r.get("row_name"))
        groups[(code_family(code), key, markers)].append(r)

    tpl_groups: dict[tuple[str, str, str], list[TemplateRowSlot]] = defaultdict(list)
    ordered_slots = sorted(slots.slots, key=lambda x: x.order)
    for s in ordered_slots:
        tpl_groups[(s.family, s.key, s.markers)].append(s)
    referenced = referenced_row_codes(db_rows)

    bound_ids: set[int] = set()

    def _bind(slot: TemplateRowSlot, row: dict, how: str) -> None:
        copy = dict(row)
        copy["db_row_code"] = row.get("row_code")
        copy["row_code"] = slot.code
        result.by_slot[slot.slot_key] = copy
        result.by_code.setdefault(slot.code, copy)
        result.methods[how] += 1
        bound_ids.add(id(row))

    for gkey, tslots in tpl_groups.items():
        cands = groups.get(gkey, [])
        if not gkey[1]:
            # 模板行没有行名（合成模板 / 行名不在本行）：无从按行名核对，只能按编码取（旧口径）
            for s in tslots:
                same = by_code.get(s.code)
                if same is not None:
                    _bind(s, same, "code_unlabeled")
                else:
                    result.unresolved.append(s)
            continue
        if not cands:
            for s in tslots:
                result.unresolved.append(s)
                same = by_code.get(s.code)
                if same is not None:
                    result.conflicts.append((s, str(same.get("row_name") or "")))
            continue
        if len(tslots) == 1:
            s = tslots[0]
            best = _rank(cands, s.code, referenced)
            if best.get("row_code") == s.code:
                how = "code"
            elif len(cands) == 1:
                how = "name"
            else:
                how = "name_ranked"
            _bind(s, best, how)
            continue
        # 模板同名多处：先同编码配对（同一编码出现多处时，前后编码连号的「本位」优先），余下按顺序配对
        assigned: dict[str, dict] = {}
        used: set[int] = set()
        by_priority = sorted(tslots, key=lambda x: (-_code_neighbor_score(x, ordered_slots), x.order))
        for s in by_priority:
            hit = next((r for r in cands if r.get("row_code") == s.code and id(r) not in used), None)
            if hit is not None:
                assigned[s.slot_key] = hit
                used.add(id(hit))
        free = [r for r in cands if id(r) not in used]
        for s in tslots:
            if s.slot_key in assigned:
                _bind(s, assigned[s.slot_key], "code")
            elif free:
                _bind(s, free.pop(0), "order")
            else:
                result.unresolved.append(s)

    represented = {k for k in tpl_groups} | set(slots.formula_keys)
    for gkey, cands in groups.items():
        if gkey in represented:
            continue
        for r in cands:
            if id(r) not in bound_ids and _nonzero(r):
                result.unbound_nonzero.append(r)
    return result


def best_rows_by_key(rows: Iterable[dict]) -> dict[tuple[str, str, str], dict]:
    """(编码族, 行名键, 标记) → 该名下最可信的库内行（与绑定同一排序口径）。"""
    db_rows = [r for r in rows if isinstance(r, dict) and r.get("row_code")]
    referenced = referenced_row_codes(db_rows)
    groups: dict[tuple[str, str, str], list[dict]] = defaultdict(list)
    for r in db_rows:
        key, markers = normalize_row_label(r.get("row_name"))
        if key:
            groups[(code_family(r["row_code"]), key, markers)].append(r)
    return {k: _rank(v, "", referenced) for k, v in groups.items()}


_CURRENT_HEADS = ("期末余额", "本期金额", "本年金额", "期末数", "本期发生额", "本年累计数", "本期数")
_PRIOR_HEADS = ("期初余额", "上期金额", "上年金额", "期初数", "上年年末余额", "上期发生额", "上年同期累计数", "上期数")


def period_columns(ws) -> tuple[int | None, int | None]:
    """按表头文字（前 6 行）找本期 / 上期金额列；找不到返回 None。"""
    cur = prior = None
    for row in ws.iter_rows(min_row=1, max_row=6):
        for c in row:
            if not isinstance(getattr(c, "value", None), str):
                continue
            t = re.sub(r"\s", "", c.value)
            if cur is None and t in _CURRENT_HEADS:
                cur = c.column
            elif prior is None and t in _PRIOR_HEADS:
                prior = c.column
    return cur, prior


def bind_index(index: dict[str, dict], slots: TemplateSlots) -> BindingResult:
    """``{row_code: row}`` 形态的索引（母公司个别列）走同一绑定。"""
    rows = []
    for code, r in index.items():
        row = dict(r)
        row.setdefault("row_code", code)
        rows.append(row)
    return bind_rows_to_template(rows, slots)

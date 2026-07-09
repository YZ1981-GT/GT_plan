"""
Phase 0.1: 读取 N3 递延所得税负债.xlsx 全部6 sheet结构
产出: .kiro/specs/n3-deferred-tax-liabilities/n3_structure_summary.json

提取：sheet名/列头/行数/公式单元格/合并区域/数据类型
确认：审定表N3-1(24×14,78公式)/明细表N3-2(31×14,14公式)/审计程序表N3A/调整分录N3-3/附注
"""
import json
import os
import re
from pathlib import Path
from openpyxl import load_workbook
from openpyxl.utils import get_column_letter

XLSX_PATH = (
    r"d:\GT_plan\数据\致同通用审计程序及底稿模板（2025年修订）"
    r"\1.致同审计程序及底稿模板（2025年）"
    r"\4.风险应对-实质性程序（D-N）"
    r"\N 税金循环"
    r"\N3 递延所得税负债.xlsx"
)

OUTPUT_PATH = (
    r"d:\GT_plan\.kiro\specs\n3-deferred-tax-liabilities\n3_structure_summary.json"
)


def extract_sheet_structure(ws):
    """Extract structure from a single worksheet."""
    info = {
        "title": ws.title,
        "dimensions": ws.dimensions,
        "max_row": ws.max_row,
        "max_column": ws.max_column,
        "size": f"{ws.max_row}×{ws.max_column}",
        "merged_cells": [str(m) for m in ws.merged_cells.ranges],
        "merged_count": len(ws.merged_cells.ranges),
        "headers_row1": [],
        "headers_row2": [],
        "headers_row3": [],
        "formula_cells": [],
        "formula_count": 0,
        "business_formula_cells": [],
        "business_formula_count": 0,
        "header_ref_count": 0,
        "formula_patterns": [],
        "sample_data_rows": [],
        "data_types": {},
    }

    # Extract row 1, 2, 3 headers
    for col in range(1, min(ws.max_column + 1, 51)):  # cap at 50 cols
        col_letter = get_column_letter(col)
        for row_idx, key in [(1, "headers_row1"), (2, "headers_row2"), (3, "headers_row3")]:
            cell = ws.cell(row=row_idx, column=col)
            val = cell.value
            if val is not None:
                info[key].append({"col": col_letter, "value": str(val)[:120]})

    # Scan for formula cells and data types
    type_counts = {}
    for row in ws.iter_rows(min_row=1, max_row=ws.max_row,
                            min_col=1, max_col=ws.max_column):
        for cell in row:
            # Track data types
            dt = cell.data_type  # 's'=string, 'n'=numeric, 'f'=formula, 'b'=bool, 'd'=date
            type_counts[dt] = type_counts.get(dt, 0) + 1

            if cell.value and isinstance(cell.value, str) and cell.value.startswith("="):
                info["formula_cells"].append({
                    "cell": f"{get_column_letter(cell.column)}{cell.row}",
                    "formula": cell.value[:150],  # truncate long formulas
                })

    info["formula_count"] = len(info["formula_cells"])
    info["data_types"] = type_counts

    # Separate business formulas from header references (底稿目录!)
    header_refs = [fc for fc in info["formula_cells"] if "底稿目录!" in fc["formula"]]
    business_formulas = [fc for fc in info["formula_cells"] if "底稿目录!" not in fc["formula"]]
    info["header_ref_count"] = len(header_refs)
    info["business_formula_cells"] = business_formulas
    info["business_formula_count"] = len(business_formulas)

    # Distinct formula patterns (normalize cell references)
    patterns = set()
    for fc in business_formulas:
        formula = fc["formula"]
        pattern = re.sub(r"[A-Z]+\d+", "REF", formula)
        patterns.add(pattern)
    info["formula_patterns"] = sorted(patterns)

    # Count unique formula rows (business only)
    biz_rows = set()
    for fc in business_formulas:
        row_num = re.search(r"\d+", fc["cell"]).group()
        biz_rows.add(row_num)
    info["business_formula_rows"] = len(biz_rows)

    # Sample first 8 data rows (row 3~10) for context
    for row_idx in range(3, min(11, ws.max_row + 1)):
        row_data = {}
        for col in range(1, min(ws.max_column + 1, 31)):
            cell = ws.cell(row=row_idx, column=col)
            if cell.value is not None:
                col_letter = get_column_letter(col)
                val = cell.value
                if isinstance(val, str) and val.startswith("="):
                    row_data[col_letter] = f"[FORMULA] {val[:100]}"
                else:
                    row_data[col_letter] = str(val)[:100]
        if row_data:
            info["sample_data_rows"].append({"row": row_idx, "data": row_data})

    return info


def main():
    print(f"Loading workbook: {XLSX_PATH}")
    if not Path(XLSX_PATH).exists():
        print(f"❌ File not found: {XLSX_PATH}")
        return

    wb = load_workbook(XLSX_PATH, data_only=False)  # data_only=False to see formulas

    print(f"Found {len(wb.sheetnames)} sheets: {wb.sheetnames}")

    sheets_info = []
    for ws_name in wb.sheetnames:
        ws = wb[ws_name]
        print(f"  Processing: {ws_name} ({ws.max_row}×{ws.max_column})")
        info = extract_sheet_structure(ws)
        sheets_info.append(info)

    # Build summary
    summary = {
        "source_file": os.path.basename(XLSX_PATH),
        "total_sheets": len(wb.sheetnames),
        "sheet_names": wb.sheetnames,
        "key_sheets": {
            "N3-1_审定表": None,
            "N3-2_明细表": None,
            "N3A_审计程序表": None,
            "N3-3_调整分录": None,
            "附注": None,
        },
        "sheets": sheets_info,
        "total_formulas": sum(s["formula_count"] for s in sheets_info),
        "total_business_formulas": sum(s["business_formula_count"] for s in sheets_info),
        "validation": {
            "expected_sheet_count": 6,
            "actual_sheet_count": len(wb.sheetnames),
            "N3-1_size": None,
            "N3-1_formula_count": None,
            "N3-2_size": None,
            "N3-2_formula_count": None,
        },
    }

    # Identify key sheets by pattern matching
    for s in sheets_info:
        title = s["title"]
        if "N3-1" in title or ("审定" in title and "N3" in title):
            summary["key_sheets"]["N3-1_审定表"] = {
                "title": title,
                "size": s["size"],
                "total_formula_count": s["formula_count"],
                "business_formula_count": s["business_formula_count"],
                "header_ref_count": s["header_ref_count"],
                "formula_patterns": s["formula_patterns"],
                "business_formula_rows": s.get("business_formula_rows", 0),
            }
            summary["validation"]["N3-1_size"] = s["size"]
            summary["validation"]["N3-1_formula_count"] = s["business_formula_count"]
        elif "N3-2" in title or ("明细" in title and "N3" in title):
            summary["key_sheets"]["N3-2_明细表"] = {
                "title": title,
                "size": s["size"],
                "total_formula_count": s["formula_count"],
                "business_formula_count": s["business_formula_count"],
                "header_ref_count": s["header_ref_count"],
                "formula_patterns": s["formula_patterns"],
                "business_formula_rows": s.get("business_formula_rows", 0),
            }
            summary["validation"]["N3-2_size"] = s["size"]
            summary["validation"]["N3-2_formula_count"] = s["business_formula_count"]
        elif "N3A" in title or ("程序" in title and "N3" in title):
            summary["key_sheets"]["N3A_审计程序表"] = {
                "title": title,
                "size": s["size"],
                "total_formula_count": s["formula_count"],
                "business_formula_count": s["business_formula_count"],
                "header_ref_count": s["header_ref_count"],
                "formula_patterns": s["formula_patterns"],
                "business_formula_rows": s.get("business_formula_rows", 0),
            }
        elif "N3-3" in title or ("调整" in title and "N3" in title):
            summary["key_sheets"]["N3-3_调整分录"] = {
                "title": title,
                "size": s["size"],
                "total_formula_count": s["formula_count"],
                "business_formula_count": s["business_formula_count"],
                "header_ref_count": s["header_ref_count"],
                "formula_patterns": s["formula_patterns"],
                "business_formula_rows": s.get("business_formula_rows", 0),
            }
        elif "附注" in title or "disclosure" in title.lower():
            summary["key_sheets"]["附注"] = {
                "title": title,
                "size": s["size"],
                "total_formula_count": s["formula_count"],
                "business_formula_count": s["business_formula_count"],
                "header_ref_count": s["header_ref_count"],
                "formula_patterns": s["formula_patterns"],
                "business_formula_rows": s.get("business_formula_rows", 0),
            }
        # Also check for 底稿目录 (Tab Index)
        if "底稿目录" in title or "目录" in title:
            summary["key_sheets"]["底稿目录"] = {
                "title": title,
                "size": s["size"],
                "total_formula_count": s["formula_count"],
                "business_formula_count": s["business_formula_count"],
                "header_ref_count": s["header_ref_count"],
                "formula_patterns": s["formula_patterns"],
                "business_formula_rows": s.get("business_formula_rows", 0),
            }

    # Print validation summary
    print("\n=== VALIDATION ===")
    v = summary["validation"]
    print(f"  Sheet数量: {v['actual_sheet_count']} (预期6)")
    sheets_match = "✓" if v["actual_sheet_count"] == 6 else "≠"
    print(f"    {sheets_match}")

    print(f"\n  N3-1 审定表:")
    print(f"    尺寸: {v['N3-1_size']} (预期24×14)")
    print(f"    业务公式数: {v['N3-1_formula_count']} (预期78)")

    print(f"\n  N3-2 明细表:")
    print(f"    尺寸: {v['N3-2_size']} (预期31×14)")
    print(f"    业务公式数: {v['N3-2_formula_count']} (预期14)")

    print(f"\n  总公式: {summary['total_formulas']}")
    print(f"  总业务公式: {summary['total_business_formulas']}")

    # Formula analysis per key sheet
    print("\n=== 公式分析 ===")
    for key, sheet_info in summary["key_sheets"].items():
        if sheet_info:
            print(f"  {key}: {sheet_info['business_formula_rows']}行有业务公式, "
                  f"{len(sheet_info['formula_patterns'])}种模式, "
                  f"{sheet_info['business_formula_count']}个公式cell")

    # Write output
    os.makedirs(os.path.dirname(OUTPUT_PATH), exist_ok=True)
    with open(OUTPUT_PATH, "w", encoding="utf-8") as f:
        json.dump(summary, f, ensure_ascii=False, indent=2)

    print(f"\n✅ Output written to: {OUTPUT_PATH}")


if __name__ == "__main__":
    main()

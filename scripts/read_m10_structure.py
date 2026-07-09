"""
Phase 0.1: 读取 M10 其他权益工具.xlsx 全部sheet结构
产出: .kiro/specs/m10-other-equity-instruments/m10_structure_summary.json
"""
import json
import os
import re
from pathlib import Path
from openpyxl import load_workbook
from openpyxl.utils import get_column_letter

XLSX_PATH = (
    r"d:\GT_plan\数据\致同通用审计程序及底稿模板（2025年修订）"
    r"\1.致同审计程序及底稿模板（2025年）"
    r"\4.风险应对-实质性程序（D-N）"
    r"\M 权益循环"
    r"\M10 其他权益工具.xlsx"
)

OUTPUT_PATH = (
    r"d:\GT_plan\.kiro\specs\m10-other-equity-instruments\m10_structure_summary.json"
)


def extract_sheet_structure(ws):
    """Extract structure from a single worksheet."""
    info = {
        "title": ws.title,
        "dimensions": ws.dimensions,
        "max_row": ws.max_row,
        "max_column": ws.max_column,
        "size": f"{ws.max_row}×{ws.max_column}",
        "merged_cells": [str(m) for m in ws.merged_cells.ranges],
        "merged_count": len(ws.merged_cells.ranges),
        "headers_row1": [],
        "headers_row2": [],
        "formula_cells": [],
        "formula_count": 0,
        "business_formula_cells": [],
        "business_formula_count": 0,
        "header_ref_count": 0,
        "formula_patterns": [],
        "sample_data_rows": [],
    }

    # Extract row 1 and row 2 headers
    for col in range(1, min(ws.max_column + 1, 51)):  # cap at 50 cols
        cell_r1 = ws.cell(row=1, column=col)
        cell_r2 = ws.cell(row=2, column=col)
        val_r1 = cell_r1.value
        val_r2 = cell_r2.value
        col_letter = get_column_letter(col)

        if val_r1 is not None:
            info["headers_row1"].append({"col": col_letter, "value": str(val_r1)})
        if val_r2 is not None:
            info["headers_row2"].append({"col": col_letter, "value": str(val_r2)})

    # Scan for formula cells
    for row in ws.iter_rows(min_row=1, max_row=ws.max_row,
                            min_col=1, max_col=ws.max_column):
        for cell in row:
            if cell.value and isinstance(cell.value, str) and cell.value.startswith("="):
                info["formula_cells"].append({
                    "cell": f"{get_column_letter(cell.column)}{cell.row}",
                    "formula": cell.value[:120],  # truncate long formulas
                })

    info["formula_count"] = len(info["formula_cells"])

    # Separate business formulas from header references (底稿目录!)
    header_refs = [fc for fc in info["formula_cells"] if "底稿目录!" in fc["formula"]]
    business_formulas = [fc for fc in info["formula_cells"] if "底稿目录!" not in fc["formula"]]
    info["header_ref_count"] = len(header_refs)
    info["business_formula_cells"] = business_formulas
    info["business_formula_count"] = len(business_formulas)

    # Distinct formula patterns (normalize cell references)
    patterns = set()
    for fc in business_formulas:
        formula = fc["formula"]
        pattern = re.sub(r"[A-Z]+\d+", "REF", formula)
        patterns.add(pattern)
    info["formula_patterns"] = sorted(patterns)

    # Count unique formula rows (business only)
    biz_rows = set()
    for fc in business_formulas:
        row_num = re.search(r"\d+", fc["cell"]).group()
        biz_rows.add(row_num)
    info["business_formula_rows"] = len(biz_rows)

    # Sample first 5 data rows (row 3~7) for context
    for row_idx in range(3, min(8, ws.max_row + 1)):
        row_data = {}
        for col in range(1, min(ws.max_column + 1, 31)):
            cell = ws.cell(row=row_idx, column=col)
            if cell.value is not None:
                col_letter = get_column_letter(col)
                val = cell.value
                if isinstance(val, str) and val.startswith("="):
                    row_data[col_letter] = f"[FORMULA] {val[:80]}"
                else:
                    row_data[col_letter] = str(val)[:80]
        if row_data:
            info["sample_data_rows"].append({"row": row_idx, "data": row_data})

    return info


def identify_skip_sheets(sheets_info):
    """Identify Q10A修订前 sheet to skip."""
    skip_sheets = []
    for s in sheets_info:
        title = s["title"]
        if "修订前" in title or "Q10A" in title:
            skip_sheets.append(title)
    return skip_sheets


def main():
    print(f"Loading workbook: {XLSX_PATH}")
    wb = load_workbook(XLSX_PATH, data_only=False)  # data_only=False to see formulas

    print(f"Found {len(wb.sheetnames)} sheets: {wb.sheetnames}")

    sheets_info = []
    for ws_name in wb.sheetnames:
        ws = wb[ws_name]
        print(f"  Processing: {ws_name} ({ws.max_row}×{ws.max_column})")
        info = extract_sheet_structure(ws)
        sheets_info.append(info)

    skip_sheets = identify_skip_sheets(sheets_info)

    # Build summary
    summary = {
        "source_file": os.path.basename(XLSX_PATH),
        "total_sheets": len(wb.sheetnames),
        "sheet_names": wb.sheetnames,
        "skip_sheets": skip_sheets,
        "active_sheets": [s for s in wb.sheetnames if s not in skip_sheets],
        "active_sheet_count": len(wb.sheetnames) - len(skip_sheets),
        "key_sheets": {
            "M10-1_审定表": None,
            "M10-2_明细表": None,
            "M10-4_区分检查": None,
        },
        "sheets": sheets_info,
        "total_formulas": sum(s["formula_count"] for s in sheets_info),
        "validation": {
            "M10-1_formula_count": None,
            "M10-2_size": None,
            "M10-2_formula_count": None,
            "M10-4_size": None,
        },
    }

    # Identify key sheets by pattern matching
    for s in sheets_info:
        title = s["title"]
        if "M10-1" in title or ("审定" in title and "M10" in title):
            summary["key_sheets"]["M10-1_审定表"] = {
                "title": title,
                "size": s["size"],
                "total_formula_count": s["formula_count"],
                "business_formula_count": s["business_formula_count"],
                "header_ref_count": s["header_ref_count"],
                "formula_patterns": s["formula_patterns"],
                "business_formula_rows": s.get("business_formula_rows", 0),
            }
            summary["validation"]["M10-1_formula_count"] = s["business_formula_count"]
        elif "M10-2" in title or ("明细" in title and "M10" in title):
            summary["key_sheets"]["M10-2_明细表"] = {
                "title": title,
                "size": s["size"],
                "total_formula_count": s["formula_count"],
                "business_formula_count": s["business_formula_count"],
                "header_ref_count": s["header_ref_count"],
                "formula_patterns": s["formula_patterns"],
                "business_formula_rows": s.get("business_formula_rows", 0),
            }
            summary["validation"]["M10-2_size"] = s["size"]
            summary["validation"]["M10-2_formula_count"] = s["business_formula_count"]
        elif "M10-4" in title or ("区分" in title and "M10" in title):
            summary["key_sheets"]["M10-4_区分检查"] = {
                "title": title,
                "size": s["size"],
                "total_formula_count": s["formula_count"],
                "business_formula_count": s["business_formula_count"],
                "header_ref_count": s["header_ref_count"],
                "formula_patterns": s["formula_patterns"],
            }
            summary["validation"]["M10-4_size"] = s["size"]

    # Print validation summary
    print("\n=== VALIDATION ===")
    v = summary["validation"]
    print(f"  M10-1 业务公式数: {v['M10-1_formula_count']} (预期29)")
    print(f"  M10-2 尺寸: {v['M10-2_size']} (预期45×30) ✓" if v['M10-2_size'] == '45×30' else f"  M10-2 尺寸: {v['M10-2_size']} (预期45×30) ✗")
    print(f"  M10-2 业务公式数: {v['M10-2_formula_count']} (预期13)")
    print(f"  M10-4 尺寸: {v['M10-4_size']} (预期64×8) ✓" if v['M10-4_size'] == '64×8' else f"  M10-4 尺寸: {v['M10-4_size']} (预期64×8) ✗")
    print(f"  跳过sheets: {skip_sheets}")
    print(f"  有效sheet数: {summary['active_sheet_count']}")
    print(f"  总公式数: {summary['total_formulas']} (业务公式: {sum(s['business_formula_count'] for s in sheets_info)})")

    # Note about formula count discrepancy
    print("\n=== 公式数说明 ===")
    print("  spec预期的'29公式'/'13公式'指的是不同公式行(row-level unique patterns)")
    print("  实际每行公式会跨多列复制，总cell数更多")
    m10_1 = summary["key_sheets"]["M10-1_审定表"]
    m10_2 = summary["key_sheets"]["M10-2_明细表"]
    if m10_1:
        print(f"  M10-1: {m10_1['business_formula_rows']}行有业务公式, {len(m10_1['formula_patterns'])}种模式, {m10_1['business_formula_count']}个公式cell")
    if m10_2:
        print(f"  M10-2: {m10_2['business_formula_rows']}行有业务公式, {len(m10_2['formula_patterns'])}种模式, {m10_2['business_formula_count']}个公式cell")

    # Write output
    os.makedirs(os.path.dirname(OUTPUT_PATH), exist_ok=True)
    with open(OUTPUT_PATH, "w", encoding="utf-8") as f:
        json.dump(summary, f, ensure_ascii=False, indent=2)

    print(f"\n✅ Output written to: {OUTPUT_PATH}")


if __name__ == "__main__":
    main()

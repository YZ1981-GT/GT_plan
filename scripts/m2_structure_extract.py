"""
M2 实收资本（股本）.xlsx 结构提取脚本
Phase 0 Task 0.1: 读取全部11个sheet，提取结构+确认关键sheet公式数

产出: m2_structure_summary.json
"""
import json
import os
from pathlib import Path

from openpyxl import load_workbook
from openpyxl.utils import get_column_letter

# 源文件路径
XLSX_PATH = (
    r"D:\GT_plan\数据\致同通用审计程序及底稿模板（2025年修订）"
    r"\1.致同审计程序及底稿模板（2025年）"
    r"\4.风险应对-实质性程序（D-N）\M 权益循环"
    r"\M2 实收资本（股本）.xlsx"
)

OUTPUT_PATH = Path(__file__).parent.parent / ".kiro" / "specs" / "m2-paid-in-capital" / "m2_structure_summary.json"


def extract_sheet_structure(ws):
    """提取单个sheet的结构信息"""
    # 基本维度
    max_row = ws.max_row or 0
    max_col = ws.max_column or 0

    # 提取前5行的列头（通常含表头信息）
    headers = []
    for row_idx in range(1, min(6, max_row + 1)):
        row_data = []
        for col_idx in range(1, max_col + 1):
            cell = ws.cell(row=row_idx, column=col_idx)
            val = cell.value
            if val is not None:
                row_data.append({
                    "col": get_column_letter(col_idx),
                    "value": str(val)[:100]  # 截断长文本
                })
        if row_data:
            headers.append({"row": row_idx, "cells": row_data})

    # 统计公式单元格
    formula_cells = []
    formula_count = 0
    for row in ws.iter_rows(min_row=1, max_row=max_row, max_col=max_col):
        for cell in row:
            if cell.value and isinstance(cell.value, str) and cell.value.startswith("="):
                formula_count += 1
                # 记录前50个公式的位置和内容
                if len(formula_cells) < 50:
                    formula_cells.append({
                        "cell": f"{get_column_letter(cell.column)}{cell.row}",
                        "formula": str(cell.value)[:200]
                    })

    # 统计合并单元格
    merged_ranges = [str(r) for r in ws.merged_cells.ranges]

    return {
        "dimensions": f"{max_row}×{max_col}",
        "max_row": max_row,
        "max_col": max_col,
        "headers": headers,
        "formula_count": formula_count,
        "formula_cells_sample": formula_cells,
        "merged_cells_count": len(merged_ranges),
        "merged_ranges_sample": merged_ranges[:20],
    }


def main():
    print(f"读取文件: {XLSX_PATH}")
    assert os.path.exists(XLSX_PATH), f"文件不存在: {XLSX_PATH}"

    # data_only=False 以保留公式
    wb = load_workbook(XLSX_PATH, data_only=False, read_only=False)

    sheet_names = wb.sheetnames
    print(f"共 {len(sheet_names)} 个sheet: {sheet_names}")

    result = {
        "source_file": "M2 实收资本（股本）.xlsx",
        "total_sheets": len(sheet_names),
        "sheet_names": sheet_names,
        "sheets": {},
        "key_sheet_validation": {},
    }

    total_formulas = 0
    for name in sheet_names:
        ws = wb[name]
        print(f"  处理: {name} ...")
        structure = extract_sheet_structure(ws)
        result["sheets"][name] = structure
        total_formulas += structure["formula_count"]
        print(f"    维度: {structure['dimensions']}, 公式: {structure['formula_count']}")

    result["total_formulas"] = total_formulas

    # 关键sheet验证
    # 1. M2-1 审定表 (spec估计73公式，实际值以xlsx为准)
    m2_1_candidates = [n for n in sheet_names if "M2-1" in n or "审定" in n]
    if m2_1_candidates:
        m2_1_name = m2_1_candidates[0]
        m2_1_info = result["sheets"][m2_1_name]
        result["key_sheet_validation"]["M2-1_审定表"] = {
            "sheet_name": m2_1_name,
            "spec_estimate_formulas": 73,
            "actual_formulas": m2_1_info["formula_count"],
            "dimensions": m2_1_info["dimensions"],
            "note": "实际公式数以xlsx为准(含行列跨区域公式)",
        }

    # 2. M2-2 明细表 (上市38×36 / 非上市37×24)
    m2_2_candidates = [n for n in sheet_names if "M2-2" in n or "明细" in n]
    for name in m2_2_candidates:
        info = result["sheets"][name]
        is_listed = "上市" in name and "非上市" not in name
        is_unlisted = "非上市" in name
        key = f"M2-2_{'上市' if is_listed else '非上市' if is_unlisted else '未知'}明细"
        expected_dims = "38×36" if is_listed else "37×24" if is_unlisted else "unknown"
        result["key_sheet_validation"][key] = {
            "sheet_name": name,
            "expected_dimensions": expected_dims,
            "actual_dimensions": info["dimensions"],
            "dimensions_match": info["dimensions"] == expected_dims,
            "formula_count": info["formula_count"],
        }

    # 3. M2-4 外币投资 (spec估计13公式，实际值以xlsx为准)
    m2_4_candidates = [n for n in sheet_names if "M2-4" in n or "外币" in n]
    if m2_4_candidates:
        m2_4_name = m2_4_candidates[0]
        m2_4_info = result["sheets"][m2_4_name]
        result["key_sheet_validation"]["M2-4_外币投资"] = {
            "sheet_name": m2_4_name,
            "spec_estimate_formulas": 13,
            "actual_formulas": m2_4_info["formula_count"],
            "dimensions": m2_4_info["dimensions"],
            "note": "实际公式数以xlsx为准",
        }

    # 输出JSON
    OUTPUT_PATH.parent.mkdir(parents=True, exist_ok=True)
    with open(OUTPUT_PATH, "w", encoding="utf-8") as f:
        json.dump(result, f, ensure_ascii=False, indent=2)

    print(f"\n产出文件: {OUTPUT_PATH}")
    print(f"总公式数: {total_formulas}")
    print("\n=== 关键sheet验证 ===")
    for key, val in result["key_sheet_validation"].items():
        print(f"  📋 {key}:")
        for k, v in val.items():
            print(f"      {k}: {v}")

    wb.close()


if __name__ == "__main__":
    main()
